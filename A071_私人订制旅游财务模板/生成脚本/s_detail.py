# -*- coding: utf-8 -*-
"""查看表（绿）【订单详情】一单的结算单（A4 竖向能打印，可以给客户签字）。
   黄格 C3：订单号（下拉＝订单列表，没结算的；已结算的直接打订单号；选的是「号｜客户｜出行 m/d」取「｜」前面；
            空＝订单登记里最后一个有效订单：单_有效 分 60 块数，MATCH 找最后一块，块里再 MATCH 找最后一个）。
   抬头：P_公司名称＋「订单结算单」；订单信息卡（左）、金额卡（右）：MATCH(键, 单_键, 0) 找到第几条，再 INDEX 单_xxx。
   下面是一块「活动区」（隐藏列先算出每一行是什么，打印区域跟着伸缩）：
     ① 收款明细：第一行定金（订单登记填的），下面收支登记里这单 归类＝订单收款/订单退款 的每一笔（按日期排），合计＋核对
        （已付、日期≤截止的 收入－退出 ＝ 单_已收）；最多 60 行。
     ② 成本明细：收支登记里这单 归类＝订单成本 的每一笔（按日期排），金额＝支出－收入，其中未付；合计＝单_实际成本（核对）；最多 80 行。
     ③ 利润小结：订单总金额（已取消＝留下的钱）－实际成本＝实际利润；预计利润；差额；订单提成。签字栏。
   排序：本表隐藏列（_收 每条一格，折成 160 行×50 列的方块，两块：①、②）只放这单的 收_排序键，SMALL 取第 k 条，MOD 10000 还原第几条。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'I'
R_SEL, R_HINT = 3, 4
R_HEAD, R_SUB, R_CH, R_C0 = 6, 7, 8, 9
NC = 12                                  # 信息卡、金额卡各 12 行
R0 = R_C0 + NC + 1                       # 活动区第一行（22）
N1, N2 = 60, 80                          # ① ② 最多显示几行（① 含定金那行）
NREG = 160                               # 活动区行数（够 19＋N1＋N2）
R1 = R0 + NREG - 1
assert 19 + N1 + N2 <= NREG

# ── 隐藏列 ──
SL, SV = 'AA', 'AB'                      # 标量：说明、值（第 3 行起）
SC = ['选', '号', '有效数', '末块', '最后', '找到', 'n', '键', '定金', '定金日期', 'dep', 'c1', 'c2', 'm1', 'm2', 'E1', 'E2',
      '末行', '收入合计', '退出合计', '到账合计', '成本合计', '未付合计', '核对1', '核对2', '出行码', '应发', '有预计', '截止']
SR = {k: 3 + i for i, k in enumerate(SC)}
BC_, BI_ = 'AC', 'AD'                    # 找最后一个有效订单：AC 每 50 条一块的累计个数（60 块）；AD 最后那块里逐条累计（50 条）
QB = 50
NQB = N_ORD // QB
assert N_ORD % QB == 0
HK, HT, HX = 'AE', 'AF', 'AG'            # 活动区：第 k 行、这一行是什么、_收 第几条
HIN, HOUT, HGOT, HCOST, HUNP = 'AH', 'AI', 'AJ', 'AK', 'AL'     # ① 收入、退出、已到账；② 金额、其中未付
BLK = 160                                # _收 8000 条折成 160 行×50 列
assert N_CASH % BLK == 0 and BLK <= NREG
NBK = N_CASH // BLK
BN = 'AN'                                # 方块行号 1～160
K1 = [CL(CI('AO') + j) for j in range(NBK)]                 # ① 的键（AO..CL）
K2 = [CL(CI(K1[-1]) + 1 + j) for j in range(NBK)]           # ② 的键（CM..EJ）
KR1 = f'${K1[0]}${R0}:${K1[-1]}${R0 + BLK - 1}'
KR2 = f'${K2[0]}${R0}:${K2[-1]}${R0 + BLK - 1}'

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_BIG = Font(name=YH, sz=16, bold=True, color='FF000000')
F_SUBH = Font(name=YH, sz=11, bold=True, color='FF000000')
F_GREY_I = Font(name=YH, sz=10, italic=True, color='FF9E9E9E')
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_LBL = fill('FFD9E1F2')
GREEN_H = 'FF70AD47'
MONEY_B = '#,##0.00;[Red]-#,##0.00;"-"'


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def UL(n):
    """签字线（下划线文字：Excel、WPS、LibreOffice 显示一样）"""
    return '"' + '_' * n + '"'


def _sw(cd, m):
    return '=' + _chain([(f'{cd}="{k}"', v) for k, v in m.items()])


TIP = ('💡 一单一张结算单，可以直接打印（A4 竖向）给客户、老板签字。黄格选订单号：下拉只列没结算的单（新的在上面）；'
       '已结算的单下拉里没有，直接打订单号；空着＝订单登记里最后一个订单。上面两张卡是订单信息和金额（已收、还没收按首页截止日算），'
       '下面 ① 收款明细（定金＋收支登记里这单的收尾款/分期、退款）、② 成本明细（这单的每一笔成本，含未付的）、③ 利润小结。'
       '灰色斜体＝还没定下来的数：已出行没结算的实际利润是「暂算」，没结算的订单提成是「预估」。要改哪一笔，按「登记行号」去【收支登记】改。')


def build(wb, ctx=None):
    ws = wb[SH_DETAIL]
    widths(ws, {'A': 12, 'B': 13, 'C': 12, 'D': 27, 'E': 13.5, 'F': 13.5, 'G': 10, 'H': 10, 'I': 8, 'J': 10})
    title(ws, '订单详情（结算单）', LAST, C_VIEW, TIP)
    home_link(ws, 'J1')

    # ── 选择格 ──
    selector(ws, 'A3', '订单号', 'C3', None, '=订单列表',
             prompt='从下拉选；已结算的单下拉里没有，直接打订单号；空着＝订单登记里最后一个订单')
    ws.merge_cells('A3:B3')
    ws['B3'].border = BD
    ws.merge_cells('C3:D3')
    ws['D3'].border = BD
    ws.merge_cells(f'E3:{LAST}3')
    N = S('n')
    put(ws, 'E3', (f'=IF({S("选")}="","空着＝订单登记里最后一个订单"&IF({N}>0,"："&INDEX(单_标签,{N}),"（还没有订单）"),'
                   f'IF({N}>0,"✓ "&INDEX(单_标签,{N}),IF({S("找到")}>0,"✗ 订单登记里这一行没填全（不算订单）",'
                   f'"✗ 订单登记里没有这个订单号")))'), F_NOTE, align=AL, border=False)
    ws.merge_cells(f'A{R_HINT}:{LAST}{R_HINT}')
    put(ws, f'A{R_HINT}', '提示：下拉只列没结算的订单；已结算的单下拉里没有，直接打订单号（比如 26001）。灰色斜体＝暂算、预估（还没结算）。',
        F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add('E3', FormulaRule(formula=['LEFT($E$3,1)="✗"'], font=F_WARN))
    ws.row_dimensions[R_SEL].height = 24
    ws.row_dimensions[5].height = 8

    # ── 标量 ──
    sel, no = S('选'), S('号')
    nokey = '"#"&' + no
    rngBC = f'${BC_}$3:${BC_}${2 + NQB}'
    rngBI = f'${BI_}$3:${BI_}${2 + QB}'
    sc = {
        '选': '=TRIM(C3&"")',
        '号': f'=IF({sel}="","",IFERROR(TRIM(LEFT({sel},FIND("｜",{sel})-1)),{sel}))',
        '有效数': f'=${BC_}${2 + NQB}',
        '末块': f'=IF({S("有效数")}=0,0,IFERROR(MATCH({S("有效数")},{rngBC},0),0))',
        '最后': (f'=IF({S("末块")}=0,0,({S("末块")}-1)*{QB}+IFERROR(MATCH({S("有效数")},{rngBI},0),0))'),
        '找到': f'=IF({no}="",0,IFERROR(MATCH({esc(nokey)},单_键,0),0))',
        'n': (f'=IF({no}="",{S("最后")},IF({S("找到")}=0,0,IF(INDEX(单_有效,{S("找到")})=1,{S("找到")},0)))'),
        '键': f'=IF({N}=0,"@@没有这单@@",INDEX(单_键,{N}))',
        '定金': f'=IF({N}=0,0,INDEX(单_定金,{N}))',
        '定金日期': f'=IF({N}=0,0,INDEX(单_定金日期,{N}))',
        'dep': f'=IF({S("定金")}<>0,1,0)',
        'c1': f'=COUNT({KR1})',
        'c2': f'=COUNT({KR2})',
        'm1': f'={S("dep")}+MIN({S("c1")},{N1}-{S("dep")})',
        'm2': f'=MIN({S("c2")},{N2})',
        'E1': f'=2+{S("m1")}',
        'E2': f'={S("E1")}+4+{S("m2")}',
        '末行': f'={R0 - 1}+{S("E2")}+13',
        '收入合计': f'=ROUND(SUM(${HIN}${R0}:${HIN}${R1}),2)',
        '退出合计': f'=ROUND(SUM(${HOUT}${R0}:${HOUT}${R1}),2)',
        '到账合计': f'=ROUND(SUM(${HGOT}${R0}:${HGOT}${R1}),2)',
        '成本合计': f'=ROUND(SUM(${HCOST}${R0}:${HCOST}${R1}),2)',
        '未付合计': f'=ROUND(SUM(${HUNP}${R0}:${HUNP}${R1}),2)',
        '核对1': f'=IF({N}=0,1,IF(ROUND({S("到账合计")}-INDEX(单_已收,{N}),2)=0,1,0))',
        '核对2': (f'=IF({N}=0,1,IF(AND(ROUND({S("成本合计")}-INDEX(单_实际成本,{N}),2)=0,'
                f'ROUND({S("未付合计")}-INDEX(单_未付成本,{N}),2)=0),1,0))'),
        '出行码': f'=IF({N}=0,-1,INDEX(单_出行码,{N}))',
        '应发': f'=IF({N}=0,0,INDEX(单_应发,{N}))',
        '有预计': f'=IF({N}=0,0,INDEX(单_有预计,{N}))',
        '截止': '=P_截止',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP
    # 最后一个有效订单：先数每 50 条一块的累计个数，找最后一块；块里再逐条累计
    for j in range(1, NQB + 1):
        ws[f'{BC_}{2 + j}'] = f'=COUNTIF(INDEX(单_有效,1):INDEX(单_有效,{QB * j}),1)'
        ws[f'{BC_}{2 + j}'].font = F_HELP
    for i in range(1, QB + 1):
        ws[f'{BI_}{2 + i}'] = (f'=IF({S("末块")}=0,0,COUNTIF(INDEX(单_有效,1):INDEX(单_有效,({S("末块")}-1)*{QB}+{i}),1))')
        ws[f'{BI_}{2 + i}'].font = F_HELP

    cx, cy, ccode = S('出行码'), S('应发'), S('有预计')
    g = lambda k: f'INDEX(单_{k},{N})'
    V = lambda k: f'=IF({N}=0,"",{g(k)})'
    Vd = lambda k: f'=IF({N}=0,"",IF({g(k)}=0,"",{g(k)}))'

    # ── 抬头 ──
    ws.merge_cells(f'A{R_HEAD}:{LAST}{R_HEAD}')
    put(ws, f'A{R_HEAD}', '=IF(P_公司名称="","",P_公司名称&"　")&"订单结算单"', F_BIG, align=AC, border=False)
    ws.row_dimensions[R_HEAD].height = 32
    ws.merge_cells(f'A{R_SUB}:E{R_SUB}')
    put(ws, f'A{R_SUB}', f'=IF({N}=0,"订单号：（没有这单）","订单号："&{g("订单号")}&"　　客户："&{g("客户名字")})',
        F_SUBH, align=AL, border=False)
    ws.merge_cells(f'F{R_SUB}:{LAST}{R_SUB}')
    put(ws, f'F{R_SUB}', '="截至 "&TEXT(P_截止,"yyyy/mm/dd")', F_TXT, align=AR, border=False)
    ws.row_dimensions[R_SUB].height = 22

    # ── 两张卡 ──
    ws.merge_cells(f'A{R_CH}:C{R_CH}')
    put(ws, f'A{R_CH}', '订单信息', F_SEC, fill(C_VIEW), align=AC)
    for c in 'BC':
        ws[f'{c}{R_CH}'].border = BD
    ws.merge_cells(f'D{R_CH}:{LAST}{R_CH}')
    put(ws, f'D{R_CH}', '金额（单位：元）', F_SEC, fill(C_VIEW), align=AC)
    for c in 'EFGHI':
        ws[f'{c}{R_CH}'].border = BD
    ws.row_dimensions[R_CH].height = 22
    trv = f'IF({g("出行日期")}>0,{g("出行日期")}-P_截止,0)'
    left = [
        ('订单号', V('订单号'), None),
        ('客户名字', V('客户名字'), None),
        ('联系电话', V('联系电话'), None),
        ('线路/目的地', V('线路'), None),
        ('人数', Vd('人数'), INT),
        ('销售名字', V('销售'), None),
        ('预订日期', Vd('预订日期'), DATE),
        ('出行日期', f'=IF({N}=0,"",IF({g("出行日期")}=0,"没填",{g("出行日期")}))', DATE),
        ('出行状态', (f'=IF({N}=0,"",{g("出行状态")}&IF({cx}=1,"（出行 "&{g("出行天数")}&" 天，还没结算）",'
                     f'IF(AND({cx}=0,{trv}>0),"（还有 "&{trv}&" 天出发）","")))'), None),
        ('取消日期', f'=IF({N}=0,"",IF({g("取消日期")}=0,IF({g("取消填了")}=1,"填错了（按已取消算）",""),{g("取消日期")}))', DATE),
        ('结算日期', Vd('结算日期'), DATE),
        ('登记行号', f'=IF({N}=0,"","订单登记第 "&{g("录入行")}&" 行")', None),
    ]
    dd = S('定金日期')
    has_est = f'OR({ccode}=1,{cx}=3)'
    right = [
        ('订单总金额', V('订单总金额'), f'=IF({cx}=3,"已取消：留下的钱＝已收－应退未退＝"&TEXT({g("有效金额")},"#,##0.00"),"")'),
        ('预计成本', f'=IF({N}=0,"",IF({ccode}=1,{g("预计成本")},"没填"))',
         f'=IF(AND({N}>0,{ccode}=0),"没填预计成本：预计利润按 0，不算进未出行预估总利润","")'),
        ('预计利润', f'=IF({N}=0,"",IF({has_est},{g("预计利润")},"没填预计成本"))',
         f'=IF({N}=0,"",IF({cx}=3,"已取消：＝实际利润","＝订单总金额－预计成本"))'),
        ('定金金额', V('定金'), (f'=IF(OR({N}=0,{S("定金")}=0),"",IF({dd}=0,"没填日期，不算收到","收款日期 "&TEXT({dd},"yyyy/mm/dd")'
                               f'&"　账户 "&{g("收款账户")}&IF({dd}>P_截止,"（晚于截止日，暂不算收到）","")))')),
        ('已收', V('已收'), f'=IF({N}=0,"","定金＋收尾款/分期－已退（已付、截止日以前）")'),
        ('还没收', V('还没收'), (f'=IF({N}=0,"",IF({cx}=3,"已取消，不用再收",IF(AND({cx}=0,{g("还没收")}>0.005,{trv}>0,'
                              f'{trv}<=P_尾款天数),"⚠ "&{trv}&" 天后出发，尾款还没收齐","")))')),
        ('尾款收否', V('尾款收否'), None),
        ('实际成本', V('实际成本'), (f'=IF({N}=0,"",IF({g("未付成本")}<>0,"其中未付 "&TEXT({g("未付成本")},"#,##0.00")&"（欠供应商）",'
                                f'"含未付的；供应商退回、返佣冲减"))')),
        ('实际利润', f'=IF({N}=0,"",IF({cx}=0,"未出行",{g("实际利润")}))',
         (f'=IF({N}=0,"",CHOOSE({cx}+1,"未出行：成本还没记全，先不算","（暂算）已出行还没结算","已结算",'
          f'"已取消：留下的钱－实际成本"))')),
        ('提成比例', V('提成比例'), f'=IF({N}=0,"","提成基数："&{g("提成基数")})'),
        ('订单提成', V('订单提成'), (f'=IF({N}=0,"",IF({cy}=1,"应发（已结算，算在 "&TEXT({g("结算日期")},"yyyy年m月")&"）",'
                                f'IF({cx}=3,"已取消，没有提成","（预估）还没结算，结算后才定")))')),
        ('应退未退', V('退款未付'), f'=IF({N}=0,"",IF({g("退款未付")}<>0,"答应退客户还没退（收支登记里付款情况＝未付）",""))'),
    ]
    rows_of = {}
    for i in range(NC):
        r = R_C0 + i
        lb, f, fm = left[i]
        put(ws, f'A{r}', lb, F_KPI_L, F_LBL, align=AC)
        ws.merge_cells(f'B{r}:C{r}')
        put(ws, f'B{r}', f, F_AUTOB, None, fm, AL)
        ws[f'C{r}'].border = BD
        lb, f, note = right[i]
        rows_of[lb] = r
        put(ws, f'D{r}', lb, F_KPI_L, F_LBL, align=AC)
        fm = PCT if lb == '提成比例' else MONEY
        put(ws, f'E{r}', f, F_AUTOB, None, fm, AR)
        ws.merge_cells(f'F{r}:{LAST}{r}')
        put(ws, f'F{r}', note, F_NOTE, None, align=AL)
        for c in 'GHI':
            ws[f'{c}{r}'].border = BD
        ws.row_dimensions[r].height = 20
    ws[f'E{rows_of["尾款收否"]}'].alignment = AR
    ws.conditional_formatting.add(f'E{rows_of["尾款收否"]}', FormulaRule(
        formula=[f'OR($E${rows_of["尾款收否"]}="未收",$E${rows_of["尾款收否"]}="部分",LEFT($E${rows_of["尾款收否"]},2)="多收")'],
        font=Font(bold=True, color='FFC65911')))
    ws.conditional_formatting.add(f'E{rows_of["实际利润"]}:F{rows_of["实际利润"]}',
                                  FormulaRule(formula=[f'{cx}=1'], font=F_GREY_I))
    ws.conditional_formatting.add(f'E{rows_of["订单提成"]}:F{rows_of["订单提成"]}',
                                  FormulaRule(formula=[f'AND({N}>0,{cy}=0,{cx}<>3)'], font=F_GREY_I))
    ws.conditional_formatting.add(f'F{R_C0}:F{R_C0 + NC - 1}', FormulaRule(formula=[f'LEFT($F{R_C0},1)="⚠"'], font=F_WARN))
    ws.conditional_formatting.add(f'E{rows_of["应退未退"]}', FormulaRule(formula=[f'N($E${rows_of["应退未退"]})<>0'], font=F_WARN))

    # ── 隐藏：排序键方块（只放这单的） ──
    key = S('键')
    ws[f'{BN}{R0 - 1}'] = '方块行号'
    ws[f'{K1[0]}{R0 - 1}'] = '①的键'
    ws[f'{K2[0]}{R0 - 1}'] = '②的键'
    for i in range(BLK):
        r = R0 + i
        ws[f'{BN}{r}'] = i + 1
        ws[f'{BN}{r}'].font = F_HELP
        for j in range(NBK):
            n = f'(${BN}{r}+{j * BLK})' if j else f'${BN}{r}'
            x = lambda k: f'INDEX(收_{k},{n})'
            ws[f'{K1[j]}{r}'] = (f'=IF({x("订单键")}<>{key},"",IF(OR({x("归类")}="订单收款",{x("归类")}="订单退款"),'
                                 f'{x("排序键")},""))')
            ws[f'{K2[j]}{r}'] = f'=IF({x("订单键")}<>{key},"",IF({x("归类")}="订单成本",{x("排序键")},""))'
            ws[f'{K1[j]}{r}'].font = ws[f'{K2[j]}{r}'].font = F_HELP

    # ── 活动区 ──
    E1, E2, dep = S('E1'), S('E2'), S('dep')
    for c, t in ((HK, 'k'), (HT, '类型'), (HX, '_收第几条'), (HIN, '①收入'), (HOUT, '①退出'), (HGOT, '①已到账'),
                 (HCOST, '②金额'), (HUNP, '②未付')):
        ws[f'{c}{R0 - 1}'] = t
        ws[f'{c}{R0 - 1}'].font = F_HELP
    fm = {'A': DATE, 'E': MONEY, 'F': MONEY, 'I': '0'}
    al = {'A': AC, 'B': AL, 'C': AL, 'D': Alignment(horizontal='left', vertical='center', shrink_to_fit=True),
          'E': AR, 'F': AR, 'G': AC, 'H': AC, 'I': AC}
    pv = f'OR({ccode}=1,{cx}=3)'
    for kk in range(NREG):
        r = R0 + kk
        k, T, X = f'${HK}{r}', f'${HT}{r}', f'${HX}{r}'
        ws[f'{HK}{r}'] = kk + 1
        ws[f'{HT}{r}'] = '=' + _chain([
            (f'{k}=1', '"S1"'), (f'{k}=2', '"H1"'),
            (f'{k}<={E1}', f'IF(AND({k}=3,{dep}=1),"DD","D1")'),
            (f'{k}={E1}+1', '"T1"'), (f'{k}={E1}+3', '"S2"'), (f'{k}={E1}+4', '"H2"'),
            (f'AND({k}>={E1}+5,{k}<={E2})', '"D2"'),
            (f'{k}={E2}+1', '"T2"'), (f'{k}={E2}+3', '"S3"'),
            (f'AND({k}>={E2}+4,{k}<={E2}+9)', f'"P"&({k}-{E2}-3)'),
            (f'{k}={E2}+11', '"G1"'), (f'{k}={E2}+13', '"G2"'),
        ])
        ws[f'{HX}{r}'] = (f'=IF({T}="D1",MOD(SMALL({KR1},{k}-2-{dep}),10000),'
                          f'IF({T}="D2",MOD(SMALL({KR2},{k}-{E1}-4),10000),0))')
        x = lambda f_: f'INDEX(收_{f_},{X})'
        ws[f'{HIN}{r}'] = f'=IF({T}="DD",{S("定金")},IF({T}="D1",{x("收入")},0))'
        ws[f'{HOUT}{r}'] = f'=IF({T}="D1",{x("支出")},0)'
        ws[f'{HGOT}{r}'] = (f'=IF({T}="DD",IF(AND({dd}>0,{dd}<=P_截止),{S("定金")},0),'
                            f'IF({T}="D1",IF(AND({x("未付")}=0,{x("日期")}<=P_截止),{x("收入")}-{x("支出")},0),0))')
        ws[f'{HCOST}{r}'] = f'=IF({T}="D2",{x("支出")}-{x("收入")},0)'
        ws[f'{HUNP}{r}'] = f'=IF({T}="D2",IF({x("未付")}=1,{x("支出")}-{x("收入")},0),0)'
        for c in (HK, HT, HX, HIN, HOUT, HGOT, HCOST, HUNP):
            ws[f'{c}{r}'].font = F_HELP
        nz = lambda v: f'IF({v}=0,"",{v})'
        st1 = (f'IF({x("未付")}=1,IF({x("归类")}="订单退款","应退未退","未付"),IF({x("日期")}>P_截止,"截止日以后",'
               f'IF({x("收入")}-{x("支出")}>=0,"已收","已退")))')
        chk1 = (f'IF({N}=0,"",IF({dep}+{S("c1")}>{N1},"⚠ 共 "&({dep}+{S("c1")})&" 笔，只列了前 {N1} 笔",'
                f'IF({S("核对1")}=1,"✓ 已到账 "&TEXT({S("到账合计")},"#,##0.00")&"＝已收",'
                f'"✗ 已到账 "&TEXT({S("到账合计")},"#,##0.00")&"≠已收 "&TEXT(INDEX(单_已收,{N}),"#,##0.00"))))')
        chk2 = (f'IF({N}=0,"",IF({S("c2")}>{N2},"⚠ 共 "&{S("c2")}&" 笔，只列了前 {N2} 笔",'
                f'IF({S("核对2")}=1,"✓ ＝实际成本"&IF({S("未付合计")}<>0,"（其中未付 "&TEXT({S("未付合计")},"#,##0.00")&"）",""),'
                f'"✗ 跟实际成本 "&TEXT(INDEX(单_实际成本,{N}),"#,##0.00")&" 对不上")))')
        na = f'IF({N}=0,"",'
        disp = {
            'A': {'S1': '"① 收款明细"', 'H1': '"日期"', 'DD': f'IF({dd}>0,{dd},"")', 'D1': x('日期'),
                  'T1': f'IF({dep}+{S("c1")}>{N1},"前 {N1} 笔合计","合计")',          # 超出显示行数：合计只是列出来的这些
                  'S2': '"② 成本明细"', 'H2': '"日期"', 'D2': x('日期'), 'T2': f'IF({S("c2")}>{N2},"前 {N2} 笔合计","合计")',
                  'S3': '"③ 利润小结"',
                  'G1': '"经办人："', 'G2': '"日　期："'},
            'B': {'H1': '"项目"', 'DD': '"定金"', 'D1': x('类别'), 'T1': f'"共 "&({dep}+{S("c1")})&" 笔"',
                  'H2': '"收支类别"', 'D2': x('类别'), 'T2': f'"共 "&{S("c2")}&" 笔"', 'G1': UL(12), 'G2': UL(12)},
            'C': {'H1': '"往来对象"', 'DD': g('客户名字'), 'D1': x('往来对象'), 'H2': '"往来对象"', 'D2': x('往来对象'),
                  'G1': '"审　核："', 'G2': '"日　期："'},
            'D': {'S1': '"定金＋收尾款/分期－退款（含应退未退的）"', 'H1': '"摘要"',
                  'DD': (f'"订单登记第 "&{g("录入行")}&" 行的定金"&IF({dd}=0,"（没填日期，不算收到）",'
                         f'IF({dd}>P_截止,"（晚于截止日，暂不算）",""))'),
                  'D1': x('摘要'), 'T1': chk1,
                  'S2': '"这单的地接、机票、酒店……（含未付的；金额＝支出－收入）"', 'H2': '"摘要"', 'D2': x('摘要'), 'T2': chk2,
                  'S3': '"截至 "&TEXT(P_截止,"yyyy/mm/dd")',
                  'P1': f'IF({cx}=3,"留下的钱（已收－应退未退）","订单总金额")',
                  'P2': '"－ 实际成本（含未付）"', 'P3': '"＝ 实际利润"', 'P4': '"预计利润"', 'P5': '"差额（实际－预计）"',
                  'P6': f'{na}"订单提成（"&TEXT({g("提成比例")},"0.0%")&" × "&{g("提成基数")}&"）")',
                  'G1': UL(26), 'G2': UL(26)},
            'E': {'H1': '"收入"', 'DD': S('定金'), 'D1': nz(x('收入')), 'T1': S('收入合计'),
                  'H2': '"金额"', 'D2': f'${HCOST}{r}', 'T2': S('成本合计'),
                  'P1': f'{na}{g("有效金额")})', 'P2': f'{na}{g("实际成本")})', 'P3': f'{na}{g("实际利润")})',
                  'P4': f'{na}IF({pv},{g("预计利润")},"没填预计成本"))',
                  'P5': f'{na}IF({pv},ROUND({g("实际利润")}-{g("预计利润")},2),""))',
                  'P6': f'{na}{g("订单提成")})'},
            'F': {'H1': '"退出"', 'D1': nz(x('支出')), 'T1': S('退出合计'),
                  'H2': '"其中未付"', 'D2': nz(f'${HUNP}{r}'), 'T2': S('未付合计'),
                  'P3': f'{na}CHOOSE({cx}+1,"未出行（暂算）","暂算","","已取消"))',
                  'P4': f'{na}IF({cx}=3,"＝实际利润",""))',
                  'P5': f'{na}IF(AND({cx}<=1,{pv}),"暂算",""))',
                  'P6': f'{na}IF({cy}=1,"应发",IF({cx}=3,"已取消","预估")))',
                  'G1': '"客户确认："', 'G2': '"日　期："'},
            'G': {'H1': '"账户"', 'DD': g('收款账户'), 'D1': x('账户'), 'H2': '"账户"', 'D2': x('账户'), 'G1': UL(10), 'G2': UL(10)},
            'H': {'H1': '"付款情况"', 'DD': f'IF({dd}=0,"没有日期",IF({dd}>P_截止,"截止日以后","已收"))', 'D1': st1,
                  'H2': '"已付/未付"', 'D2': f'IF({x("未付")}=1,"未付","已付")', 'G1': UL(10), 'G2': UL(10)},
            'I': {'H1': '"登记行号"', 'D1': x('录入行'), 'H2': '"登记行号"', 'D2': x('录入行'), 'G1': UL(8), 'G2': UL(8)},
        }
        for col, m in disp.items():
            c = ws[f'{col}{r}']
            c.value = _sw(T, m)
            c.font = F_TXT
            c.alignment = al[col]
            if col in fm:
                c.number_format = fm[col]

    # ── 活动区的样子（条件格式） ──
    reg = f'A{R0}:{LAST}{R1}'
    t0 = f'${HT}{R0}'
    F_WB = Font(bold=True, color='FFFFFFFF')
    cf = ws.conditional_formatting
    cf.add(f'D{R0}:D{R1}', FormulaRule(formula=[f'AND(OR({t0}="T1",{t0}="T2"),OR(LEFT($D{R0},1)="✗",LEFT($D{R0},1)="⚠"))'],
                                       font=Font(bold=True, color='FFC00000')))
    cf.add(f'H{R0}:H{R1}', FormulaRule(formula=[f'OR($H{R0}="应退未退",$H{R0}="未付")'], font=Font(bold=True, color='FFC65911')))
    cf.add(f'E{R0}:F{R0 + NREG - 1}', FormulaRule(
        formula=[f'OR(AND(OR({t0}="P3",{t0}="P5"),{cx}<=1),AND({t0}="P6",{cy}=0,{cx}<>3))'], font=Font(italic=True, color='FF9E9E9E')))
    cf.add(reg, FormulaRule(formula=[f'OR({t0}="S1",{t0}="S2",{t0}="S3")'], fill=cf_fill(C_VIEW), font=F_WB, stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'OR({t0}="H1",{t0}="H2")'], fill=cf_fill(GREEN_H), font=F_WB, border=CF_BD,
                            stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'OR({t0}="T1",{t0}="T2")'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD,
                            stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'OR({t0}="DD",{t0}="D1",{t0}="D2")'], border=CF_BD, stopIfTrue=True))
    cf.add(f'D{R0}:F{R1}', FormulaRule(formula=[f'{t0}="P3"'], font=Font(bold=True), border=CF_BD, fill=cf_fill('FFFCE4D6')))
    cf.add(f'D{R0}:F{R1}', FormulaRule(formula=[f'LEFT({t0},1)="P"'], border=CF_BD))
    cf.add(f'A{R0}:{LAST}{R1}', FormulaRule(formula=[f'OR({t0}="G1",{t0}="G2")'], font=Font(bold=True)))
    for r in range(R0, R1 + 1):
        ws.row_dimensions[r].height = 18

    hide(ws, *[CL(i) for i in range(CI('K'), CI(K2[-1]) + 1)])
    ws.freeze_panes = 'A5'
    print_setup(ws, f'{R_HEAD}:{R_SUB}', landscape=False)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    ws.print_options.horizontalCentered = True
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A${R_HEAD}:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("末行")})')
    return ws
