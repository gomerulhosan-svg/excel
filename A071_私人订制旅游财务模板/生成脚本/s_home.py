# -*- coding: utf-8 -*-
"""【首页】：截止日期（全书默认）、本月/本年关键数（用户原话的名字）、全部预计订单未出行预估总利润、该跟进的、各账户余额、提醒、各表入口、常见情况怎么记。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *
import s_check

NAV = [
    ('平时录（蓝）', C_IN, [(SH_ORD, '接单登记一行：客户、日期、金额、定金、销售'),
                       (SH_CASH, '定金以外的钱：尾款、退款、订单成本、工资、提成、房租……')]),
    ('查看（绿，自动）', C_VIEW, [(SH_LIST, '每个订单详细清单（按日期、销售、状态筛）'), (SH_DETAIL, '选一个订单：收款、成本、利润明细（结算单）'),
                           (SH_TODO, '未出行订单、预估总利润、没收齐的、欠人家的'), (SH_MONTH, '1～12 月：预订、回款、支出、工资、提成、利润'),
                           (SH_PAY, '每人工资、提成（应发、已发、未发）'), (SH_FLOW, '现金、银行、微信、支付宝日记账')]),
    ('基础（灰）', C_MASTER, [(SH_BASE, '公司名称、账户、收支类别、人员（提成比例）')]),
    ('校验', C_CHK, [(SH_CHK, '哪里录错了、该跟进什么')]),
]

HOWTO = [
    ('接单、收定金', '【订单登记】一行：订单号、客户名字、预订日期、出行日期、销售名字、订单总金额、预计成本、定金金额（定金不是当天收的填定金收款日期）'),
    ('收尾款、分期款', '【收支登记】一行：类别选「收尾款/分期款」，选订单号，填收入金额和账户（一次收齐、分几次收都这么记）'),
    ('付地接、机票、酒店等成本', '【收支登记】一行：类别选地接费、机票、酒店住宿……，选订单号，填支出金额。开票税费、收款手续费、渠道返佣也这么记'),
    ('成本先欠着（地接回来再结）', '【收支登记】照记一行，付款情况选「未付」：算进这单的成本，不算钱出去。付了以后清掉「未付」、日期改成付款那天；'
                         '只付了一部分：原行改成还欠的数，另起一行记付了的'),
    ('客户加项目、减项目、部分退团、给折让', '直接改【订单登记】的「订单总金额」，备注写原价和原因；多收的钱记「收尾款/分期款」，退的钱记「退客户款」'),
    ('客户取消', '【订单登记】填取消日期；退给客户的钱在【收支登记】记「退客户款」（答应退还没退的，付款情况选「未付」）；扣损给地接的记订单成本'),
    ('结团、成本对完账、尾款收齐', '【订单登记】填结算日期：实际利润、订单提成就定下来了，提成算在结算那个月'),
    ('发工资、发提成', '【收支登记】类别选「工资」或「销售提成」，往来对象选人员；工资和提成一起转的也拆成两行'),
    ('房租、推广、办公等日常开销', '【收支登记】选对应类别，不用选订单号'),
    ('供应商返佣', '按单返的：类别「供应商返佣（按单）」，选订单号，填收入栏（冲这单成本）；分不到订单的季度返佣：「供应商返佣（季度）」，算其他收入'),
    ('微信提现到对公、现金存银行', '【收支登记】记两行「内部转账」：转出账户那行填支出，转入账户那行填收入，金额一样；不算收支'),
    ('老板投钱、拿钱，借款还款', '【收支登记】类别选「股东投入」「股东取出」「借款」「还借款」：进出资金，不算收入支出'),
]


def build(wb, ctx=None):
    ws = wb[SH_HOME]
    widths(ws, {'A': 2, 'B': 18, 'C': 15, 'D': 15, 'E': 3, 'F': 18, 'G': 15, 'H': 15, 'I': 3, 'J': 18, 'K': 15, 'L': 15})
    title(ws, '私 人 订 制 旅 游 财 务', 'L', C_HOME,
          '💡 平时只录两张蓝色的表：【订单登记】（接单时一单一行，收了定金填定金）和【收支登记】（之后这单的尾款、退款、成本，还有工资、提成、房租……每笔一行）；'
          '订单清单、订单详情、未出行订单、月度汇总、工资提成、资金流水都是自动的。两张登记表、基础资料都可以插行、删行。')
    ws['A1'].value = '=IF(P_公司名称="","",P_公司名称&" · ")&"私人订制旅游财务"'
    put(ws, 'B4', '截止日期', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, HOME_END, None, F_SEL, FILL_SEL, DATE, AC)
    dv_date(ws, HOME_END)
    put(ws, 'D4', '="现在按 "&TEXT(P_截止,"yyyy-mm-dd")&" 算（空着＝最后一笔 "&TEXT(P_最后日期,"yyyy-mm-dd")&"）；建账日 "&TEXT(P_建账日,"yyyy-mm-dd")',
        Font(name=YH, sz=11, bold=True, color='FF1F3864'), align=AL, border=False)
    ws.merge_cells('D4:L4')
    put(ws, 'B5', (f'=IF(AND(ISNUMBER({HOME_END}),P_截止<P_最后日期),"⚠ 截止日期比最新录入（"&TEXT(P_最后日期,"yyyy-mm-dd")&"）早，'
                   f'后面录的都没算进来——是不是忘了清空？","截止日期决定：哪些单算已出行、已收和资金算到哪天、「本月」是哪个月。平时空着就行（＝最后一笔的日期）；'
                   f'想按今天算，在黄格里输入 =TODAY()。")'), F_NOTE, align=AL, border=False)
    ws.merge_cells('B5:L5')
    ws.conditional_formatting.add('B5', FormulaRule(formula=['LEFT(B5,1)="⚠"'], font=F_RED))

    # ① 本月 / 本年
    Y0 = 'YEAR(P_截止)*100+1'
    MON = lambda nm: f'{nm},P_截止年月'
    YR = lambda nm: f'{nm},">="&({Y0}),{nm},"<="&P_截止年月'
    dep = lambda cr: f'SUMIFS(单_定金,单_有效,1,{cr("单_定金年月")},单_定金日期,">="&P_建账日,单_定金日期,"<="&P_截止)'
    k = lambda kind, cr: f'SUMIFS(收_净额,收_归类,"{kind}",{cr("收_年月")})'

    def rows(cr):
        out = '+'.join(f'-{k(x, cr)}' for x in s_check.OUT_KINDS)
        settle = f'SUMIFS(单_实际利润,单_应发,1,{cr("单_结算年月")})'
        cancel = f'SUMIFS(单_实际利润,单_有效,1,单_出行码,3,{cr("单_取消年月")})'
        due = f'SUMIFS(单_订单提成,单_应发,1,{cr("单_结算年月")})'
        return [
            ('月预订金额', f'SUMIFS(单_订单总金额,单_有效,1,{cr("单_预订年月")})'),
            ('月回款金额', f'{dep(cr)}+{k("订单收款", cr)}+{k("订单退款", cr)}'),
            ('月支出金额', out),
            ('月工资', f'-{k("工资", cr)}'),
            ('销售提成（已发）', f'-{k("销售提成", cr)}'),
            ('销售提成（应发）', due),
            ('结算利润（实际）', settle),
            ('经营利润（估）', f'{settle}+{cancel}+{k("其他收入", cr)}+{k("工资", cr)}-{due}+{k("日常费用", cr)}+{k("其他支出", cr)}'),
        ]
    r = 7
    section(ws, r, 'B', 'L', '① 本月、本年（本月＝截止日期那个月；本年＝那年 1 月到截止日期；点名字看月度汇总）', C_HOME)
    header(ws, r + 1, [('B', '项目'), ('C', '本月'), ('D', '本年'), ('F', '项目'), ('G', '本月'), ('H', '本年')], C_HOME, height=22)
    mrows, yrows = rows(MON), rows(YR)
    for i, ((lab, fm), (_l, fy)) in enumerate(zip(mrows, yrows)):
        c0, cm, cy = (('B', 'C', 'D'), ('F', 'G', 'H'))[i // 4]
        rr = r + 2 + i % 4
        link(put(ws, f'{c0}{rr}', lab, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), fill('FFD9E1F2'), align=AL),
             SH_MONTH)
        put(ws, f'{cm}{rr}', f'=ROUND({fm},2)', F_AUTOB, fill('FFFFFFFF'), MONEY, AR)
        put(ws, f'{cy}{rr}', f'=ROUND({fy},2)', F_AUTOB, fill('FFFFFFFF'), MONEY, AR)
        ws.row_dimensions[rr].height = 22
    put(ws, f'J{r + 2}', '口径', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, f'K{r + 2}', ('回款、支出按收到/付出的日期，只算建账日～截止日、已付的；预订按预订日期；'
                         '工资、已发提成按发的日期；应发提成、结算利润按结算日期；经营利润（估）＝结算利润＋取消单留存＋其他收入－工资－应发提成－日常费用－其他支出。'),
        F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'K{r + 2}:L{r + 5}')
    ws.merge_cells(f'J{r + 2}:J{r + 5}')

    # ② 未出行订单
    r = r + 7
    section(ws, r, 'B', 'L', '② 未出行订单（截至截止日期；点名字看明细）', C_VIEW)
    u = '单_有效,1,单_出行码,0'
    big = f'=ROUND(SUMIFS(单_预计利润,{u},单_有预计,1),2)'
    link(put(ws, f'B{r + 1}', '全部预计订单未出行预估总利润', Font(name=YH, sz=11, bold=True, color='FF0563C1', underline='single'),
             fill('FFD9E1F2'), align=ACW), SH_TODO)
    ws.merge_cells(f'B{r + 1}:C{r + 2}')
    put(ws, f'D{r + 1}', big, Font(name=YH, sz=18, bold=True, color='FFC00000'), fill('FFFFFFFF'), MONEY, AC)
    ws.merge_cells(f'D{r + 1}:D{r + 2}')
    ws.row_dimensions[r + 1].height = 22
    ws.row_dimensions[r + 2].height = 22
    smalls = [('未出行单数', f'=COUNTIFS({u})', '0'), ('订单总金额', f'=SUMIFS(单_订单总金额,{u})', MONEY),
              ('已收', f'=SUMIFS(单_已收,{u})', MONEY), ('还没收', f'=SUMIFS(单_还没收,{u})', MONEY),
              ('预计提成', f'=SUMIFS(单_订单提成,{u})', MONEY), ('预估利润扣掉提成', f'=D{r + 1}-SUMIFS(单_订单提成,{u},单_有预计,1)', MONEY)]
    for i, (lab, f, fmt) in enumerate(smalls):
        c = ['F', 'G', 'H', 'J', 'K', 'L'][i]
        put(ws, f'{c}{r + 1}', lab, F_KPI_L, fill('FFD9E1F2'), align=AC)
        put(ws, f'{c}{r + 2}', f, F_AUTOB, fill('FFFFFFFF'), fmt, AC)
    put(ws, f'B{r + 3}', (f'=IF(COUNTIFS({u},单_有预计,0)>0,"另有 "&COUNTIFS({u},单_有预计,0)&" 单（订单金额 "'
                          f'&TEXT(SUMIFS(单_订单总金额,{u},单_有预计,0),"#,##0")&"）没填预计成本，没算进预估总利润","全部未出行订单都填了预计成本 √")'),
        F_NOTE, align=AL, border=False)
    ws.merge_cells(f'B{r + 3}:L{r + 3}')

    # ③ 该跟进的
    r = r + 5
    section(ws, r, 'B', 'L', '③ 该跟进的', C_VIEW)
    soon = '单_有效,1,单_出行码,0,单_出行日期,">0",单_出行日期,"<="&(P_截止+P_尾款天数),单_还没收,">0.005"'
    went = '单_有效,1,单_出行码,">=1",单_出行码,"<=2",单_还没收,">0.005"'
    late = '单_有效,1,单_出行码,1,单_出行日期,"<"&(P_截止-P_结算天数)'
    fol = [('="快出发（"&P_尾款天数&" 天内）没收齐"', f'=COUNTIFS({soon})&" 单"', f'=SUMIFS(单_还没收,{soon})', SH_TODO),
           ('已出行还没收齐', f'=COUNTIFS({went})&" 单"', f'=SUMIFS(单_还没收,{went})', SH_TODO),
           ('已出行还没结算', '=COUNTIFS(单_有效,1,单_出行码,1)&" 单"', '=SUMIFS(单_预计利润,单_有效,1,单_出行码,1)', SH_TODO),
           ('="出行超过 "&P_结算天数&" 天没结算"', f'=COUNTIFS({late})&" 单"', f'=SUMIFS(单_预计利润,{late})', SH_TODO),
           ('欠供应商（成本未付）', '=COUNTIFS(收_有效,1,收_未付,1,收_归类,"订单成本")&" 笔"', '=SUMIFS(收_未付额,收_有效,1,收_归类,"订单成本")', SH_TODO),
           ('应退客户还没退', '=COUNTIFS(收_有效,1,收_未付,1,收_归类,"订单退款")&" 笔"', '=SUMIFS(收_未付额,收_有效,1,收_归类,"订单退款")', SH_TODO)]
    header(ws, r + 1, [('B', '事项'), ('C', '几单/几笔'), ('D', '金额'), ('F', '事项'), ('G', '几单/几笔'), ('H', '金额')], C_VIEW, height=22)
    for i, (lab, n, amt, sh) in enumerate(fol):
        c0, cn, ca = (('B', 'C', 'D'), ('F', 'G', 'H'))[i // 3]
        rr = r + 2 + i % 3
        link(put(ws, f'{c0}{rr}', lab, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), fill('FFD9E1F2'), align=AL), sh)
        put(ws, f'{cn}{rr}', n, F_AUTOB, align=AC)
        put(ws, f'{ca}{rr}', f'=ROUND({amt[1:]},2)', F_AUTOB, fmt=MONEY, align=AR)
    put(ws, f'J{r + 2}', '金额说明', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, f'K{r + 2}', '没收齐的是还没收的钱；没结算的是这些单的预计利润；欠供应商、应退客户是收支登记里「未付」的。', F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'K{r + 2}:L{r + 4}')
    ws.merge_cells(f'J{r + 2}:J{r + 4}')

    # ④ 各账户余额
    r = r + 6
    section(ws, r, 'B', 'L', '④ 各账户余额（截至截止日期；点名字看资金流水）', C_HOME)
    acc = lambda a: (f'INDEX(账户_期初余额,MATCH({esc(a)},账户_名称,0))+SUMIFS(单_定金,单_有效,1,单_收款账户,{esc(a)},'
                     f'单_定金日期,">="&P_建账日,单_定金日期,"<="&P_截止)+SUMIFS(收_净额,收_账户,{esc(a)})')
    NA = 8
    for i in range(NA):
        c0, cv = (('B', 'C'), ('F', 'G'), ('J', 'K'))[i % 3] if i < 6 else (('B', 'C'), ('F', 'G'))[i - 6]
        rr = r + 1 + (i // 3 if i < 6 else 2)
        a = f'{H_TAB}!${COMPACT["账户"][0]}${i + 2}'
        link(put(ws, f'{c0}{rr}', f'=IF({a}="","",{a})', Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'),
                 fill('FFD9E1F2'), align=AL), SH_FLOW)
        put(ws, f'{cv}{rr}', f'=IF({a}="","",ROUND({acc(a)},2))', F_AUTOB, fill('FFFFFFFF'), MONEY, AR)
    rr = r + 3
    total = (f'SUM(账户_期初余额)+SUMIFS(单_定金,单_有效,1,单_定金日期,">="&P_建账日,单_定金日期,"<="&P_截止)+SUMIFS(收_净额,收_有效,1)')
    put(ws, f'J{rr}', '全部账户合计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'K{rr}', f'=ROUND({total},2)', F_TXTB, FILL_TOT, MONEY, AR)
    put(ws, f'L{rr}', f'=IF({H_TAB}!${COMPACT["账户"][0]}$1>{NA},"（还有 "&({H_TAB}!${COMPACT["账户"][0]}$1-{NA})&" 个账户没列，看资金流水）","")',
        F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'B{r + 1}:K{r + 3}', FormulaRule(formula=[f'AND(ISNUMBER(B{r + 1}),B{r + 1}<-0.005)'], font=F_RED))

    # ⑤ 提醒
    r = r + 5
    section(ws, r, 'B', 'L', '⑤ 提醒', C_CHK)
    CR = s_check.CHECK_ROWS
    cc = lambda lab: f'N({SH_CHK}!$C${CR[lab]})'
    rem = [f'="数据校验：要改 "&{SH_CHK}!{s_check.CHK_X}&" 项、要看 "&{SH_CHK}!{s_check.CHK_W}&" 项"&'
           f'IF({SH_CHK}!{s_check.CHK_X}+{SH_CHK}!{s_check.CHK_W}=0,"，都没问题 √","  → 点这里去看")',
           f'=IF({cc(SH_ORD)}+{cc(SH_CASH)}>0,"登记表有 "&({cc(SH_ORD)}+{cc(SH_CASH)})&" 行要改（看那一行最右边的「这一行的问题」）","两张登记表都没有要改的 √")',
           f'=IF({cc("内部转账没配对（转出、转入合计不为 0）")}<>0,"内部转账有一笔只记了一边（差 "&TEXT({cc("内部转账没配对（转出、转入合计不为 0）")},"#,##0.00")&"）","")',
           f'=IF({cc("没填预计成本的单数（预估利润没算进去）")}>0,{cc("没填预计成本的单数（预估利润没算进去）")}&" 单没填预计成本（预估利润没算进去）","")',
           f'=IF({cc("账户余额是负数的账户个数")}>0,"有 "&{cc("账户余额是负数的账户个数")}&" 个账户余额是负数：多半漏记了收款，或者期初余额没填","")',
           '=IF(P_公司名称="","【基础资料】还没填公司名称（首页、结算单抬头要用）","")',
           (f'=IF({cc(f"订单登记超过 {N_ORD} 单（多出来的不算）")}+{cc(f"收支登记超过 {N_CASH} 笔（多出来的不算）")}'
            f'+{cc("基础资料某一块超了（多出来的不算）")}>0,"有表超过容量了，多出来的没算进去 → 看【数据校验】","")')]
    for i, f in enumerate(rem):
        rr = r + 1 + i
        put(ws, f'B{rr}', f, F_TXTB if i == 0 else F_TXT, align=AL, border=False)
        ws.merge_cells(f'B{rr}:L{rr}')
    link(ws[f'B{r + 1}'], SH_CHK)
    ws.conditional_formatting.add(f'B{r + 1}:B{r + len(rem)}', FormulaRule(
        formula=[f'OR(ISNUMBER(SEARCH("要改 ",B{r + 1})),ISNUMBER(SEARCH("只记了一边",B{r + 1})),ISNUMBER(SEARCH("超过容量",B{r + 1})),'
                 f'ISNUMBER(SEARCH("负数",B{r + 1})),ISNUMBER(SEARCH("还没填",B{r + 1})),ISNUMBER(SEARCH("没填预计",B{r + 1})))'],
        font=Font(name=YH, sz=10, bold=True, color='FFC00000')))

    # 导航
    r = r + len(rem) + 2
    section(ws, r, 'B', 'L', '各张表（点名字跳过去；每张表右上角有「← 回首页」）', C_HOME)
    r += 1
    col_sets = [('B', 'C', 'E'), ('F', 'G', 'L')]
    for grp, color, sheets in NAV:
        put(ws, f'B{r}', grp, F_SEC, fill(color), align=AL)
        ws.merge_cells(f'B{r}:L{r}')
        r += 1
        for j, (sh, desc) in enumerate(sheets):
            c1, c2, c3 = col_sets[j % 2]
            rr = r + j // 2
            link(put(ws, f'{c1}{rr}', sh, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), align=AL), sh)
            put(ws, f'{c2}{rr}', desc, F_NOTE, align=AL)
            ws.merge_cells(f'{c2}{rr}:{c3}{rr}')
        r += (len(sheets) + 1) // 2

    # 怎么记
    r += 1
    section(ws, r, 'B', 'L', '常见情况怎么记（记住一条：接单在订单登记填定金，之后这单的每一笔钱都在收支登记记一行、选订单号）', C_HOME)
    for i, (kk, v) in enumerate(HOWTO):
        rr = r + 1 + i
        put(ws, f'B{rr}', kk, F_TXTB, fill('FFF2F2F2'), align=ALW)
        ws.merge_cells(f'B{rr}:C{rr}')
        put(ws, f'D{rr}', v, F_TXT, align=ALW)
        ws.merge_cells(f'D{rr}:L{rr}')
        ws.row_dimensions[rr].height = 32
    r = r + len(HOWTO) + 2
    section(ws, r, 'B', 'L', '注意', C_HOME)
    notes = ['① 销售名字、账户、收支类别都从下拉选；新的先到【基础资料】加一行。订单号在收支登记里从下拉选（只列没结算的单，新的在上面），也可以直接打。',
             '② 两张登记表、基础资料哪里都可以插行、删行、排序（整行一起动）；绿色的表是自动的，有保护（没有密码），只有黄格子能改。',
             '③ 预计利润＝订单总金额－预计成本；实际利润＝订单总金额－这单记的全部成本（含未付）。已出行没结算的实际利润是「暂算」（灰字），'
             '没结算的订单提成是「预估」（灰字）；填了结算日期才定下来。',
             '④ 示例行（灰底，备注里写着「示例」）正式用时整行删掉就行。',
             f'⑤ 订单最多 {N_ORD} 单、收支最多 {N_CASH} 笔；用满了另起一本（没出行的单带过去，账户期初余额填上一本的余额）。']
    for i, t in enumerate(notes):
        rr = r + 1 + i
        put(ws, f'B{rr}', t, F_TXT, align=ALW, border=False)
        ws.merge_cells(f'B{rr}:L{rr}')
        ws.row_dimensions[rr].height = 30
    print_setup(ws, None, landscape=False)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
