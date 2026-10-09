# -*- coding: utf-8 -*-
"""查看表（绿）【工资提成】每人的工资、销售提成（应发 / 已发 / 累计未发 / 预估），所选人员的订单提成明细和收到的钱。
   黄格（第 3 行，第 4 行灰字是实际用的值）：B3 年（空＝截止日那年）、D3 月（1～12；空＝全年）、F3 人员（空＝人员清单第一个）。
   期间＝[起, 止]（YYYYMM）：选了月＝那个月，没选＝那年 1～12 月。
   口径（agent_common「口径速查」）：
   - 月工资、销售提成（已发）＝收支登记 收_净额（在期：已付、建账日～截止日）取反，归类工资 / 销售提成，收_往来对象＝姓名，收_年月在期间；
   - 销售提成（应发）、结算单数/金额/利润＝单_应发=1、单_销售＝姓名、单_结算年月在期间；
   - 累计未发提成（截至截止日）＝Σ单_订单提成（应发、结算年月 ≤ 截止年月）－Σ已发（收_已付额，归类销售提成，日期 ≤ 截止，不限建账日）；
   - 预估提成＝单_出行码 0/1（没结算、没取消）的 单_订单提成（不分期间）。
   第 5 行（固定）：期间内工资、提成发给了不在人员表里的人的金额（全部－按人员加起来的）；销售不在人员表里的订单的应发提成。
   下面是一块「活动区」（第 7 行起），三块一块接一块往下排：
   ① 每人一行（人员清单＝_表 O 列压紧清单，个数在 O1）→ 合计 → 灰行「另：不在人员表里的人」；
   ② 所选人员的订单提成：期间内结算的（应发）＋全部没结算没取消的（预估），按结算日期、再按出行日期排（没结算的排在后面）；
      表头 → 应发合计 → 预估合计 → 明细（最多 300 单）→ 共 N 单；
   ③ 所选人员期间内收到的钱（收_在期=1、归类工资/销售提成、往来对象＝他），按日期排：表头 → 合计 → 明细（最多 100 笔）→ 共 N 笔。
   区块里的长说明放不下（同一行别的格子有公式，字溢不出去），所以分区标题拆成几格短字，说明写在第 2、4、5 行。
   隐藏列：AA 说明、AB 值（第 3 行起）；AC 第 k 行、AD 这一行是什么、AE 排序键、AF 取第几条、AG 应发标记；
      AI～AV 第 t 个人的数（第 R0+t-1 行）；AX 起 ② 排序键、BE 起 ③ 排序键（折成 500 行一列的方块）。
   ② 排序键＝(结算日期或 99999)×100000＋出行日期或 99999)×10000＋n（约 1E14，取 n 用 键－INT(键/10000)×10000，不用 MOD）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'M'
R0 = 7
NP, N2, N3 = N_PER, 300, 100
NREG = 5 + NP + 6 + N2 + 4 + N3
R1 = R0 + NREG - 1
BLK = 500
SL, SV = 'AA', 'AB'
HK, HC, HS, HI, HD = 'AC', 'AD', 'AE', 'AF', 'AG'
PB = dict(名='AI', 位='AJ', 岗位='AK', 基数='AL', 比例='AM', 工资='AN', 应发='AO', 已发='AP', 合计='AQ', 未发='AR', 预估='AS',
          单数='AT', 金额='AU', 利润='AV')
K2C, K3C = 'AX', 'BE'
_PO = COMPACT['人员'][0]
PLIST = f'{H_TAB}!${_PO}$2:${_PO}${N_PER + 1}'          # 人员压紧清单（没有空行）
PCNT = f'{H_TAB}!${_PO}$1'                             # 清单里几个人

SC = ['年输入', '年', '年认出', '月输入', '月', '月认出', '起', '止', '期间', '人员', '人员在表',
      'np', 'c2', 'n2', 'c3', 'n3', 'b2', 'b3', '最后一行',
      '工资合计', '应发合计', '已发合计', '工资应发合计', '未发合计', '预估合计', '单数合计', '金额合计', '利润合计',
      '全部工资', '全部已发', '全部应发', '不在表工资', '不在表已发', '不在表应发', '不在表工资应发', '不在表发钱',
      'Ta单数', 'Ta金额', 'Ta利润', 'Ta提成', 'Tb单数', 'Tb金额', 'Tb利润', 'Tb提成', '③合计', '③工资', '③提成']
SR = {k: 3 + i for i, k in enumerate(SC)}

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
GREEN_H = 'FF70AD47'
F_WHITE_B = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_GREY = Font(name=YH, sz=9, color='FF808080')
F_GREY_I = Font(name=YH, sz=10, italic=True, color='FF9E9E9E')
F_ORANGE_B = Font(name=YH, sz=10, bold=True, color='FFC65911')
LBL = fill('FFD9E1F2')
KINDS = ['S1', 'H1', 'D1', 'T1', 'N1', 'S2', 'H2', 'Ta', 'Tb', 'D2', 'N2', 'S3', 'H3', 'T3', 'D3', 'N3']


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def per(nm):
    """期间条件（YYYYMM 在 [起, 止]）"""
    return f'{nm},">="&{S("起")},{nm},"<="&{S("止")}'


def pbr(key):
    return f'${PB[key]}${R0}:${PB[key]}${R0 + NP - 1}'


def key_block(ws, c0, n, cell):
    """排序键折成 BLK 行一列的方块：第 i 条（0 起）在 c0 右边第 i//BLK 列、第 R0+i%BLK 行。返回方块区域"""
    i0 = CI(c0)
    for i in range(n):
        c = ws[f'{CL(i0 + i // BLK)}{R0 + i % BLK}']
        c.value = '=' + cell(i + 1)
        c.font = F_HELP
    ncol = -(-n // BLK)
    return f'${c0}${R0}:${CL(i0 + ncol - 1)}${R0 + BLK - 1}'


TIP = ('💡 全自动不用填。黄格：年（空着＝截止日期那年）、月（1～12；空着＝全年）、人员（空着＝人员表第一个）。'
       '① 每人一行：月工资、已发提成＝这段时间【收支登记】里发给他的（类别归到工资 / 销售提成、往来对象选他，只算已付的）；'
       '应发提成＝这段时间结算的订单（订单登记填了结算日期）该发的提成；累计未发＝到截止日期为止应发的－已发的（正数＝还欠他的，橙色）；'
       '预估提成＝还没结算的单按预计利润算的（灰色斜体，结算以后才定）。'
       '② 所选人员的订单提成：这段时间结算的（应发，按结算日期排）＋所有还没结算的（预估，按出行日期排）；利润：结算了＝实际利润，没结算＝预计利润。'
       '③ 所选人员这段时间收到的工资、提成，一笔一行。')


def build(wb, ctx=None):
    ws = wb[SH_PAY]
    W = {'A': 10, 'B': 12, 'C': 13, 'D': 14, 'E': 12, 'F': 12, 'G': 13, 'H': 12, 'I': 12, 'J': 12, 'K': 10, 'L': 12, 'M': 12}
    widths(ws, W)
    title(ws, '工资提成（每人工资、销售提成：应发、已发、未发、预估）', LAST, C_VIEW, TIP)

    # ── 第 3 行选择格、第 4 行实际用的值 ──
    selector(ws, 'A3', '年', 'B3', None, fmt='0')
    dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2099', allow_blank=True,
                        showErrorMessage=False, showInputMessage=True, promptTitle='提示', prompt='填年份，比如 2026；空着＝截止日期那年')
    ws.add_data_validation(dv)
    dv.add('B3')
    selector(ws, 'C3', '月', 'D3', None, '"' + ','.join(str(m) for m in range(1, 13)) + '"', fmt='0',
             prompt='选 1～12 月；空着＝全年')
    selector(ws, 'E3', '人员', 'F3', None, '=人员列表', prompt='选一个人看他的订单提成明细和收到的钱；空着＝人员表第一个')
    ws.merge_cells('G3:L3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 26

    pp = S('人员')
    yv, mv = S('年输入'), S('月输入')
    ok_y = f'OR(AND({yv}>=2000,{yv}<=2099),AND({yv}>=36526,{yv}<=73050))'
    np_, n2, n3, b2, b3 = S('np'), S('n2'), S('n3'), S('b2'), S('b3')
    sm = lambda k: f'ROUND(SUM({pbr(k)}),2)'
    pe = esc(pp)
    ta = f'单_有效,1,单_应发,1,单_销售,{pe},{per("单_结算年月")}'
    tb = f'单_有效,1,单_销售,{pe},单_出行码,"<=1"'
    s3 = lambda kind: f'-SUMIFS(收_净额,收_有效,1,收_归类,"{kind}",收_往来对象,{pe},{per("收_年月")})'
    g = lambda f: f'=IF({pp}="",0,ROUND({f},2))'
    sc = {
        '年输入': '=IFERROR(--TRIM(B3&""),0)',
        '年': f'=IF(AND({yv}>=2000,{yv}<=2099),INT({yv}),IF(AND({yv}>=36526,{yv}<=73050),YEAR({yv}),YEAR(P_截止)))',
        '年认出': f'=IF(TRIM(B3&"")="",1,IF({ok_y},1,0))',
        '月输入': '=IFERROR(--SUBSTITUTE(TRIM(D3&""),"月",""),0)',
        '月': f'=IF(AND({mv}>=1,{mv}<=12),INT({mv}),0)',
        '月认出': f'=IF(TRIM(D3&"")="",1,IF({S("月")}>0,1,0))',
        '起': f'={S("年")}*100+IF({S("月")}=0,1,{S("月")})',
        '止': f'={S("年")}*100+IF({S("月")}=0,12,{S("月")})',
        '期间': f'={S("年")}&"年"&IF({S("月")}=0,"全年",{S("月")}&"月")',
        '人员': f'=IF(TRIM(F3&"")="",{H_TAB}!${_PO}$2&"",TRIM(F3&""))',
        '人员在表': f'=IF({pp}="",0,IF(COUNTIF(人员_姓名,{pe})>0,1,0))',
        'np': f'=MIN(N({PCNT}),{NP})',
        'c2': '=COUNT({K2})', 'n2': f'=MIN({S("c2")},{N2})',
        'c3': '=COUNT({K3})', 'n3': f'=MIN({S("c3")},{N3})',
        'b2': f'=5+{np_}',
        'b3': f'={b2}+6+{n2}',
        '最后一行': f'={R0 - 1}+{b3}+4+{n3}',
        '工资合计': '=' + sm('工资'), '应发合计': '=' + sm('应发'), '已发合计': '=' + sm('已发'), '工资应发合计': '=' + sm('合计'),
        '未发合计': '=' + sm('未发'), '预估合计': '=' + sm('预估'), '单数合计': f'=SUM({pbr("单数")})', '金额合计': '=' + sm('金额'),
        '利润合计': '=' + sm('利润'),
        '全部工资': f'=ROUND(-SUMIFS(收_净额,收_有效,1,收_归类,"工资",{per("收_年月")}),2)',
        '全部已发': f'=ROUND(-SUMIFS(收_净额,收_有效,1,收_归类,"销售提成",{per("收_年月")}),2)',
        '全部应发': f'=ROUND(SUMIFS(单_订单提成,单_有效,1,单_应发,1,{per("单_结算年月")}),2)',
        '不在表工资': f'=ROUND({S("全部工资")}-{S("工资合计")},2)',
        '不在表已发': f'=ROUND({S("全部已发")}-{S("已发合计")},2)',
        '不在表应发': f'=ROUND({S("全部应发")}-{S("应发合计")},2)',
        '不在表工资应发': f'=ROUND({S("不在表工资")}+{S("不在表应发")},2)',
        '不在表发钱': f'=ROUND({S("不在表工资")}+{S("不在表已发")},2)',
        'Ta单数': f'=IF({pp}="",0,COUNTIFS({ta}))', 'Ta金额': g(f'SUMIFS(单_订单总金额,{ta})'),
        'Ta利润': g(f'SUMIFS(单_实际利润,{ta})'), 'Ta提成': g(f'SUMIFS(单_订单提成,{ta})'),
        'Tb单数': f'=IF({pp}="",0,COUNTIFS({tb}))', 'Tb金额': g(f'SUMIFS(单_订单总金额,{tb})'),
        'Tb利润': g(f'SUMIFS(单_本单利润,{tb})'), 'Tb提成': g(f'SUMIFS(单_订单提成,{tb})'),
        '③工资': g(s3('工资')), '③提成': g(s3('销售提成')),
        '③合计': f'=ROUND({S("③工资")}+{S("③提成")},2)',
    }

    # ── 排序键 ──
    u_ = lambda k, n: f'INDEX(单_{k},{n})'
    s_ = lambda k, n: f'INDEX(收_{k},{n})'
    K2 = key_block(ws, K2C, N_ORD, lambda n: (
        f'IF(AND({pp}<>"",{u_("有效", n)}=1,{u_("销售", n)}={pp},OR(AND({u_("应发", n)}=1,{u_("结算年月", n)}>={S("起")},'
        f'{u_("结算年月", n)}<={S("止")}),{u_("出行码", n)}<=1)),(IF({u_("应发", n)}=1,{u_("结算日期", n)},99999)*100000'
        f'+IF({u_("出行日期", n)}>0,{u_("出行日期", n)},99999))*10000+{n},"")'))
    K3 = key_block(ws, K3C, N_CASH, lambda n: (
        f'IF(AND({pp}<>"",{s_("在期", n)}=1,OR({s_("归类", n)}="工资",{s_("归类", n)}="销售提成"),{s_("往来对象", n)}={pp},'
        f'{s_("年月", n)}>={S("起")},{s_("年月", n)}<={S("止")}),{s_("排序键", n)},"")'))
    sc['c2'] = f'=COUNT({K2})'
    sc['c3'] = f'=COUNT({K3})'
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP
    for c, t in ((HK, 'k'), (HC, '这行是'), (HS, '排序键'), (HI, '第几条'), (HD, '应发'), (K2C, '②键'), (K3C, '③键'),
                 *((PB[k], k) for k in PB)):
        ws[f'{c}{R0 - 1}'] = t
        ws[f'{c}{R0 - 1}'].font = F_HELP

    # ── 第 t 个人的数（隐藏） ──
    for t in range(1, NP + 1):
        r = R0 + t - 1
        nm, pos = f'${PB["名"]}{r}', f'${PB["位"]}{r}'
        e = esc(nm)
        z = lambda f: f'=IF({nm}="",0,ROUND({f},2))'
        st = f'单_有效,1,单_应发,1,单_销售,{e},{per("单_结算年月")}'
        vals = {
            '名': f'=IF({t}>N({PCNT}),"",INDEX({PLIST},{t})&"")',
            '位': f'=IF({nm}="",0,IFERROR(MATCH({e},人员_姓名,0),0))',
            '岗位': f'=IF({pos}=0,"",INDEX(人员_岗位,{pos}))',
            '基数': f'=IF({pos}=0,"",INDEX(人员_提成基数,{pos}))',
            '比例': f'=IF({pos}=0,0,INDEX(人员_提成比例,{pos}))',
            '工资': z(f'-SUMIFS(收_净额,收_有效,1,收_归类,"工资",收_往来对象,{e},{per("收_年月")})'),
            '应发': z(f'SUMIFS(单_订单提成,{st})'),
            '已发': z(f'-SUMIFS(收_净额,收_有效,1,收_归类,"销售提成",收_往来对象,{e},{per("收_年月")})'),
            '合计': f'=ROUND(${PB["工资"]}{r}+${PB["应发"]}{r},2)',
            '未发': z(f'SUMIFS(单_订单提成,单_有效,1,单_应发,1,单_销售,{e},单_结算年月,"<="&P_截止年月)'
                    f'+SUMIFS(收_已付额,收_有效,1,收_归类,"销售提成",收_往来对象,{e},收_日期,"<="&P_截止)'),
            '预估': z(f'SUMIFS(单_订单提成,单_有效,1,单_销售,{e},单_出行码,"<=1")'),
            '单数': f'=IF({nm}="",0,COUNTIFS({st}))',
            '金额': z(f'SUMIFS(单_订单总金额,{st})'),
            '利润': z(f'SUMIFS(单_实际利润,{st})'),
        }
        for k, f in vals.items():
            ws[f'{PB[k]}{r}'] = f
            ws[f'{PB[k]}{r}'].font = F_HELP

    # ── 第 3～5 行的字 ──
    put(ws, 'G3', (f'=IF({S("年认出")}=0,"⚠ 年份没认出来（先按截止日那年）　","")'
                   f'&IF({S("月认出")}=0,"⚠ 月份要填 1～12（先按全年）　","")'
                   f'&IF(AND({pp}<>"",{S("人员在表")}=0),"⚠ 「"&{pp}&"」不在人员表里（订单、收到的钱照列）","")'),
        F_RED, align=ALW, border=False)
    put(ws, 'A4', '实际用', F_NOTE, align=AC, border=False)
    put(ws, 'B4', f'={S("年")}&" 年"', F_NOTE, align=AC, border=False)
    put(ws, 'D4', f'=IF({S("月")}=0,"全年",{S("月")}&" 月")', F_NOTE, align=AC, border=False)
    put(ws, 'F4', f'=IF({pp}="","（没有人员）",{pp})', F_NOTE, align=AC, border=False)
    ws.merge_cells(f'G4:{LAST}4')
    put(ws, 'G4', (f'="期间 "&{S("期间")}&"：工资、已发提成按发钱的日期（只算建账日～截止日、已付的）；应发提成、结算按结算日期；'
                   f'累计未发、预估提成截至 "&TEXT(P_截止,"yyyy-mm-dd")'), F_NOTE, align=ALW, border=False)
    ws.row_dimensions[4].height = 28
    ws.merge_cells(f'A5:{LAST}5')
    paid_ok, due_ok = f'ABS({S("不在表发钱")})<0.005', f'ABS({S("不在表应发")})<0.005'
    put(ws, 'A5', (f'=IF({np_}=0,"⚠ 【基础资料】的人员表是空的：先把销售、员工填进去",'
                   f'IF(AND({paid_ok},{due_ok}),"✓ 这段时间的工资、提成都发给了人员表里的人",'
                   f'"⚠ "&IF({paid_ok},"","这段时间有 "&TEXT({S("不在表发钱")},"#,##0.00")&" 元工资、提成发给了不在人员表里的人'
                   f'（【收支登记】往来对象没选人员表里的名字）")'
                   f'&IF({due_ok},"",IF({paid_ok},"","；")&"销售没填或不在人员表里的订单，这段时间应发提成 "'
                   f'&TEXT({S("不在表应发")},"#,##0.00")&" 元")&"，没算进下面每个人——见 ① 合计下面灰色那行"))'),
        F_NOTE, align=ALW, border=False)
    ws.conditional_formatting.add('A5', FormulaRule(formula=['LEFT($A$5,1)="⚠"'], font=Font(bold=True, color='FFC65911')))
    ws.row_dimensions[5].height = 28
    ws.row_dimensions[6].height = 6

    # ── 活动区 ──
    fm = {'A': '0', 'B': DATE, 'D': PCT, 'E': MONEY, 'F': MONEY, 'G': MONEY, 'H': MONEY, 'I': MONEY, 'J': MONEY, 'K': INT,
          'L': MONEY, 'M': MONEY}
    al = {'A': AC, 'B': AC, 'C': AL, 'D': AL, 'E': AR, 'F': AR, 'G': AR, 'H': AR, 'I': AR, 'J': AC, 'K': AC, 'L': AR, 'M': AR}
    for kk in range(NREG):
        r = R0 + kk
        k = f'${HK}{r}'
        cd, hs, ix = f'${HC}{r}', f'${HS}{r}', f'${HI}{r}'
        ws[f'{HK}{r}'] = kk + 1
        ws[f'{HC}{r}'] = (f'=IF({k}<=2+{np_},IF({k}=1,"S1",IF({k}=2,"H1","D1")),IF({k}=3+{np_},"T1",IF({k}=4+{np_},"N1",'
                          f'IF({k}<={b2},"",IF({k}<={b2}+4,CHOOSE({k}-{b2},"S2","H2","Ta","Tb"),IF({k}<={b2}+4+{n2},"D2",'
                          f'IF({k}={b2}+5+{n2},"N2",IF({k}<={b3},"",IF({k}<={b3}+3,CHOOSE({k}-{b3},"S3","H3","T3"),'
                          f'IF({k}<={b3}+3+{n3},"D3",IF({k}={b3}+4+{n3},"N3","")))))))))))')
        ws[f'{HS}{r}'] = f'=IF({cd}="D2",SMALL({K2},{k}-{b2}-4),IF({cd}="D3",SMALL({K3},{k}-{b3}-3),0))'
        ws[f'{HI}{r}'] = f'=IF({cd}="D1",{k}-2,IF(OR({cd}="D2",{cd}="D3"),{hs}-INT({hs}/10000)*10000,0))'
        ws[f'{HD}{r}'] = f'=IF({cd}="D2",INDEX(单_应发,{ix}),"")'
        for c in (HK, HC, HS, HI, HD):
            ws[f'{c}{r}'].font = F_HELP

        p = lambda key: f'INDEX({pbr(key)},{ix})'
        u = lambda key: f'INDEX(单_{key},{ix})'
        s = lambda key: f'INDEX(收_{key},{ix})'
        nrow = lambda c, unit: f'IF({S(c)}=0,"（没有）","共 "&{S(c)}&" {unit}")'
        capw = lambda c, cap, unit: f'IF({S(c)}>{cap},"⚠ 只列前 {cap} {unit}（合计是全部的）","")'
        D = {
            'A': {'S1': '"①"', 'H1': '"姓名"', 'D1': p('名'), 'T1': '"合计"', 'N1': '"另：不在"',
                  'S2': '"②"', 'H2': '"序号"', 'D2': f'{k}-{b2}-4', 'S3': '"③"', 'H3': '"序号"', 'D3': f'{k}-{b3}-3'},
            'B': {'S1': '"工资提成"', 'H1': '"岗位"', 'D1': p('岗位'), 'N1': '"人员表里的人"',
                  'S2': pp, 'H2': '"订单号"', 'Ta': '"应发合计"', 'Tb': '"预估合计"', 'D2': u('订单号'), 'N2': nrow('c2', '单'),
                  'S3': pp, 'H3': '"日期"', 'T3': '"合计"', 'D3': s('日期'), 'N3': nrow('c3', '笔')},
            'C': {'S1': S('期间'), 'H1': '"提成基数"', 'D1': p('基数'),
                  'S2': '"订单提成明细"', 'H2': '"客户名字"', 'Ta': f'{S("Ta单数")}&" 单"', 'Tb': f'{S("Tb单数")}&" 单"',
                  'D2': u('客户名字'), 'N2': capw('c2', N2, '单'),
                  'S3': '"收到的钱"', 'H3': '"类别"', 'D3': s('类别'), 'N3': capw('c3', N3, '笔')},
            'D': {'H1': '"提成比例"', 'D1': p('比例'),
                  'S2': S('期间'), 'H2': '"出行日期"', 'D2': f'IF({u("出行日期")}>0,TEXT({u("出行日期")},"yyyy/mm/dd"),"没填")',
                  'S3': S('期间'), 'H3': '"摘要"', 'D3': s('显示摘要')},
            'E': {'H1': '"月工资"', 'D1': p('工资'), 'T1': S('工资合计'), 'N1': S('不在表工资'),
                  'H2': '"结算日期"', 'D2': f'IF({u("应发")}=1,TEXT({u("结算日期")},"yyyy/mm/dd"),"")',
                  'H3': '"金额"', 'T3': S('③合计'), 'D3': f'-{s("净额")}'},
            'F': {'H1': '"应发提成"', 'D1': p('应发'), 'T1': S('应发合计'), 'N1': S('不在表应发'),
                  'H2': '"订单总金额"', 'Ta': S('Ta金额'), 'Tb': S('Tb金额'), 'D2': u('订单总金额'),
                  'H3': '"算作"', 'T3': '"其中工资"', 'D3': s('归类')},
            'G': {'H1': '"已发提成"', 'D1': p('已发'), 'T1': S('已发合计'), 'N1': S('不在表已发'),
                  'H2': '"实际/预计利润"', 'Ta': S('Ta利润'), 'Tb': S('Tb利润'),
                  'D2': f'IF({u("应发")}=1,{u("实际利润")},IF({u("有预计")}=1,{u("预计利润")},"没填预计"))',
                  'T3': S('③工资')},
            'H': {'H1': '"工资＋应发"', 'D1': p('合计'), 'T1': S('工资应发合计'), 'N1': S('不在表工资应发'),
                  'H2': '"提成比例"', 'D2': f'TEXT({u("提成比例")},"0.0%")&IF({u("提成基数")}="订单金额","×金额","×利润")',
                  'T3': '"提成"'},
            'I': {'H1': '"累计未发"', 'D1': p('未发'), 'T1': S('未发合计'),
                  'H2': '"订单提成"', 'Ta': S('Ta提成'), 'Tb': S('Tb提成'), 'D2': u('订单提成'), 'T3': S('③提成')},
            'J': {'H1': '"预估提成"', 'D1': p('预估'), 'T1': S('预估合计'),
                  'H2': '"状态"', 'Ta': '"应发"', 'Tb': '"预估"', 'D2': f'IF({u("应发")}=1,"应发","预估")'},
            'K': {'H1': '"结算单数"', 'D1': p('单数'), 'T1': S('单数合计'), 'H2': '"出行状态"', 'D2': u('出行状态')},
            'L': {'H1': '"结算金额"', 'D1': p('金额'), 'T1': S('金额合计')},
            'M': {'H1': '"结算利润"', 'D1': p('利润'), 'T1': S('利润合计')},
        }
        for col, m in D.items():
            x = ws[f'{col}{r}']
            x.value = '=' + _chain([(f'{cd}="{kd}"', m[kd]) for kd in KINDS if kd in m])
            x.font = F_TXT
            x.alignment = al[col]
            if col in fm:
                x.number_format = fm[col]
        ws.row_dimensions[r].height = 18

    # ── 条件格式（按隐藏列 AD 这一行是什么） ──
    full = f'A{R0}:{LAST}{R1}'
    w2, w3, w3t = f'A{R0}:K{R1}', f'A{R0}:F{R1}', f'A{R0}:I{R1}'      # ② 用到 K 列、③ 明细到 F 列、③ 合计到 I 列
    c0, d0 = f'${HC}{R0}', f'${HD}{R0}'
    is_ = lambda *ks: 'OR(' + ','.join(f'{c0}="{x}"' for x in ks) + ')'
    tot = dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD)
    rules = [
        (full, is_('S1', 'S2', 'S3'), dict(fill=cf_fill(C_VIEW), font=F_WHITE_B)),
        (full, f'AND({is_("H1", "H2", "H3")},A{R0}<>"")', dict(fill=cf_fill(GREEN_H), font=F_WHITE_B, border=CF_BD)),
        (f'J{R0}:J{R1}', f'{c0}="T1"', dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True, italic=True, color='FF9E9E9E'),
                                          border=CF_BD)),
        (f'G{R0}:I{R1}', f'{c0}="Tb"', dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True, italic=True, color='FF9E9E9E'),
                                          border=CF_BD)),
        (full, f'{c0}="T1"', tot),
        (w2, is_('Ta', 'Tb'), tot),
        (w3t, f'{c0}="T3"', tot),
        (full, f'{c0}="N1"', dict(fill=cf_fill('FFF2F2F2'), font=F_GREY, border=CF_BD)),
        (f'I{R0}:I{R1}', f'AND({c0}="D1",N(I{R0})>0.005)', dict(font=F_ORANGE_B, border=CF_BD)),
        (f'J{R0}:J{R1}', f'{c0}="D1"', dict(font=F_GREY_I, border=CF_BD)),
        (f'G{R0}:I{R1}', f'AND({c0}="D2",{d0}<>1)', dict(font=F_GREY_I, border=CF_BD)),
        (full, f'{c0}="D1"', dict(border=CF_BD)),
        (w2, f'{c0}="D2"', dict(border=CF_BD)),
        (w3, f'{c0}="D3"', dict(border=CF_BD)),
        (full, is_('N2', 'N3'), dict(font=F_GREY)),
    ]
    for rg, cond, kw in rules:
        ws.conditional_formatting.add(rg, FormulaRule(formula=[cond], stopIfTrue=True, **kw))

    hide(ws, *[CL(i) for i in range(CI(SL), CI(K3C) + -(-N_CASH // BLK))])
    ws.freeze_panes = 'A6'
    print_setup(ws, None, landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("最后一行")})')
