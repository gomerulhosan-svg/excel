# -*- coding: utf-8 -*-
"""录入表（蓝）：订单登记（一单一行）、收支登记（订单以外的钱 ＋ 订单的成本、分期追加、退款）；基础资料（灰）。
   录入表只放录入列，外加最右边几列只给人看的显示列（别的公式不引用它们，插进来的行没有这几格也不影响汇总）。"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *


def _d(v):
    if v in (None, ''):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.datetime(v.year, v.month, v.day)
    try:
        return dt.datetime.strptime(str(v)[:10], '%Y-%m-%d')
    except ValueError:
        return v


def dv_dec(ws, sqref, what='金额', lo='-999999999', hi='999999999'):
    dv = DataValidation(type='decimal', operator='between', formula1=lo, formula2=hi, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what, error=f'{what}要填数字（不要带文字、单位）')
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_len(ws, sqref, n=50):
    dv = DataValidation(type='textLength', operator='lessThanOrEqual', formula1=str(n), allow_blank=True, showErrorMessage=True,
                        errorStyle='stop', errorTitle='名称太长', error=f'名称不要超过 {n} 个字')
    ws.add_data_validation(dv)
    dv.add(sqref)


PCT1 = '0.0%;-0.0%;""'
FILL_EX = fill('FFEDEDED')
F_LBL = fill('FFD9E1F2')


def _example_grey(ws, rows, r0, last_in):
    """示例行（备注以「示例」开头）灰底，正式用时整行删掉"""
    for i, x in enumerate(rows):
        if str(x.get('备注') or '').startswith('示例'):
            for c in range(1, CI(last_in) + 1):
                ws.cell(row=r0 + i, column=c).fill = FILL_EX


def _frame(ws, last_col, hdr, heads, color, tip, W, r0, nstyle, fm, al, freeze, auto=()):
    widths(ws, W)
    title(ws, ws.title, last_col, color, tip)
    header(ws, hdr, heads, color)
    style_rows(ws, r0, r0 + nstyle - 1, [c for c, _ in heads], auto=auto, fmts=fm, aligns=al)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f'A{hdr}:{last_col}{r0 + nstyle - 1}'


def _write_rows(ws, rows, cols, r0, date_keys=('日期',)):
    for i, x in enumerate(rows):
        r = r0 + i
        for k, col in cols.items():
            v = x.get(k)
            if v in (None, ''):
                continue
            if k in date_keys:
                v = _d(v)
            ws[f'{col}{r}'] = v


def _kpi(ws, row, cells):
    """第 3 行：标签（合并两格）＋ 值"""
    for c, lbl, f, fmt in cells:
        put(ws, f'{c}{row}', lbl, F_KPI_L, F_LBL, align=AC)
        ws.merge_cells(f'{c}{row}:{CL(CI(c) + 1)}{row}')
        put(ws, f'{CL(CI(c) + 2)}{row}', f, F_AUTOB, FILL_AUTO, fmt, AR)


def _show(ws, r0, nshow, hdr, cols_formulas, fmts):
    """显示列：第 n 条＝ROW()-ROW(表头)；超出容量或没内容显示空"""
    n = f'ROW()-ROW($A${hdr})'
    for i in range(nshow):
        r = r0 + i
        for col, f in cols_formulas.items():
            ws[f'{col}{r}'] = '=IFERROR(' + f.format(n=n) + ',"")'
            ws[f'{col}{r}'].fill = FILL_AUTO
            ws[f'{col}{r}'].font = F_AUTO
            if col in fmts:
                ws[f'{col}{r}'].number_format = fmts[col]


# ───────────────────────── 订单登记 ─────────────────────────
def build_ord(ws, ctx):
    C, S = ORD_COLS, ORD_SHOW
    heads = [(C['订单号'], '订单号\n（不能重复）'), (C['客户名字'], '客户名字'), (C['预订日期'], '预订日期'), (C['出行日期'], '出行日期'),
             (C['销售'], '销售名字'), (C['订单总金额'], '订单总金额'), (C['预计成本'], '预计成本'), (C['定金金额'], '定金金额'),
             (C['联系电话'], '联系电话'), (C['线路'], '线路 / 目的地'), (C['人数'], '人数'), (C['备注'], '备注'),
             (C['定金日期'], '定金收款日期\n（空＝预订日期）'), (C['收款账户'], '定金收款账户\n（空＝默认）'),
             (C['取消日期'], '取消日期\n（取消了才填）'), (C['结算日期'], '结算日期\n（结团对完账填）'),
             (C['提成比例'], '提成比例\n（空＝按人员）'),
             (S['预计利润'], '预计利润\n（自动）'), (S['实际利润'], '实际利润\n（自动）'), (S['尾款收否'], '尾款收否\n（自动）'),
             (S['订单提成'], '订单提成\n（自动，没结算的是预估）'), (S['已收'], '已收\n（自动）'), (S['还没收'], '还没收\n（自动）'),
             (S['实际成本'], '实际成本\n（自动）'), (S['出行状态'], '出行状态\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 接单时在这里登记一行，收了定金就填定金金额。之后这一单的每一笔钱——尾款、分期、退款、地接机票酒店等成本、开票税费、手续费——'
           '都到【收支登记】记一行并选订单号，这里的已收、尾款收否、实际成本、实际利润自动算。客户加项目、减项目、部分退团、给折让：直接改「订单总金额」，备注写原价和原因。'
           '客户取消：填取消日期。结团、成本都对完账、尾款收齐：填结算日期，实际利润和提成就定下来了（提成算在结算那个月）。'
           '订单号自己编（建议数字，第 3 行有「下一个订单号」），不能重复。灰底的是示例行，正式用时整行删掉。可以插行、删行、排序（整行一起动）。')
    W = {C['订单号']: 9, C['客户名字']: 10, C['预订日期']: 11, C['出行日期']: 11, C['销售']: 8, C['订单总金额']: 12,
         C['预计成本']: 12, C['定金金额']: 11, C['联系电话']: 12, C['线路']: 20, C['人数']: 5, C['备注']: 18,
         C['定金日期']: 13, C['收款账户']: 12, C['取消日期']: 11, C['结算日期']: 12, C['提成比例']: 10,
         S['预计利润']: 12, S['实际利润']: 12, S['尾款收否']: 10, S['订单提成']: 12, S['已收']: 12, S['还没收']: 12,
         S['实际成本']: 12, S['出行状态']: 9, S['校验']: 34}
    fm = {C['预订日期']: DATE, C['出行日期']: DATE, C['定金日期']: DATE, C['取消日期']: DATE, C['结算日期']: DATE, C['人数']: INT,
          C['订单总金额']: MONEY, C['预计成本']: MONEY, C['定金金额']: MONEY, C['提成比例']: PCT1, C['订单号']: '@'}
    al = {C['客户名字']: AL, C['线路']: AL, C['备注']: AL, S['校验']: AL}
    for k in ('订单总金额', '预计成本', '定金金额'):
        al[C[k]] = AR
    rows = ctx.get('orders', [])
    nshow = N_ORD                                               # 自动列铺满容量（用户不用往下拉公式）
    _frame(ws, ORD_LAST, ORD_HDR, heads, C_IN, tip, W, ORD_R0, nshow, fm, al, f'C{ORD_R0}', auto=tuple(S.values()))
    ws.row_dimensions[ORD_HDR].height = 44
    for k in ORD_RARE:                                          # 可不填的列：表头浅色
        ws[f'{C[k]}{ORD_HDR}'].fill = fill('FF9BC2E6')
    for col in S.values():                                     # 自动列表头换个颜色
        ws[f'{col}{ORD_HDR}'].fill = fill('FF8EA9DB')
    for r in range(ORD_R0 + nshow, ORD_R0 + N_ORD):            # 订单号整列文本格式（容量以内）
        ws[f'{C["订单号"]}{r}'].number_format = '@'
    _kpi(ws, 3, [('A', '下一个订单号', '=P_下一单号', '0'),
                 ('D', '订单数（不含取消）', '=COUNTIFS(单_有效,1,单_取消,0)', '0'),
                 ('G', '订单总金额', '=SUMIFS(单_订单总金额,单_有效,1,单_取消,0)', MONEY),
                 ('J', '已收合计', '=SUMIFS(单_已收,单_有效,1)', MONEY),
                 ('M', '还没收合计', '=SUMIFS(单_还没收,单_有效,1)', MONEY)])
    put(ws, 'P3', '="已收、出行状态按截止日 "&TEXT(P_截止,"yyyy-mm-dd")&" 算（首页可改）"', F_NOTE, align=AL, border=False)
    _write_rows(ws, rows, C, ORD_R0, date_keys=('预订日期', '出行日期', '定金日期', '取消日期', '结算日期'))
    ok = 'INDEX(单_有效,{n})<>1'
    _show(ws, ORD_R0, nshow, ORD_HDR, {
        S['预计利润']: f'IF({ok},"",IF(INDEX(单_有预计,{{n}})=1,INDEX(单_预计利润,{{n}}),"没填预计成本"))',
        S['实际利润']: f'IF({ok},"",INDEX(单_实际显示,{{n}}))',
        S['尾款收否']: f'IF({ok},"",INDEX(单_尾款收否,{{n}}))',
        S['订单提成']: f'IF({ok},"",INDEX(单_订单提成,{{n}}))',
        S['已收']: f'IF({ok},"",INDEX(单_已收,{{n}}))',
        S['还没收']: f'IF({ok},"",INDEX(单_还没收,{{n}}))',
        S['实际成本']: f'IF({ok},"",INDEX(单_实际成本,{{n}}))',
        S['出行状态']: f'IF({ok},"",INDEX(单_出行状态,{{n}}))',
        S['校验']: 'INDEX(单_校验,{n})&""',
    }, {S[k]: MONEY for k in ('预计利润', '已收', '还没收', '实际成本', '实际利润', '订单提成')})
    last = ORD_R0 + nshow - 1
    ws.conditional_formatting.add(f'{S["校验"]}{ORD_R0}:{S["校验"]}{last}',
                                  FormulaRule(formula=[f'LEFT({S["校验"]}{ORD_R0},1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(f'{S["尾款收否"]}{ORD_R0}:{S["尾款收否"]}{last}',
                                  FormulaRule(formula=[f'OR({S["尾款收否"]}{ORD_R0}="未收",{S["尾款收否"]}{ORD_R0}="部分")'],
                                              font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.conditional_formatting.add(f'{S["出行状态"]}{ORD_R0}:{S["出行状态"]}{last}',
                                  FormulaRule(formula=[f'{S["出行状态"]}{ORD_R0}="未出行"'],
                                              font=Font(name=YH, sz=10, bold=True, color='FF2F75B5')))
    grey = Font(name=YH, sz=10, color='FF9E9E9E', italic=True)
    ws.conditional_formatting.add(f'{S["实际利润"]}{ORD_R0}:{S["实际利润"]}{last}',      # 已出行没结算：暂算
                                  FormulaRule(formula=[f'{S["出行状态"]}{ORD_R0}="已出行"'], font=grey))
    ws.conditional_formatting.add(f'{S["订单提成"]}{ORD_R0}:{S["订单提成"]}{last}',      # 没结算：预估
                                  FormulaRule(formula=[f'AND({S["出行状态"]}{ORD_R0}<>"已结算",{S["出行状态"]}{ORD_R0}<>"")'], font=grey))
    ws.conditional_formatting.add(f'A{ORD_R0}:{C["备注"]}{last}',
                                  FormulaRule(formula=[f'${S["出行状态"]}{ORD_R0}="已取消"'], font=Font(name=YH, sz=10, color='FF808080', strike=True)))
    rng = lambda col: f'{col}{ORD_R0}:{col}{ORD_R0 + N_ORD - 1}'
    for k in ('预订日期', '出行日期', '定金日期', '取消日期', '结算日期'):
        dv_date(ws, rng(C[k]))
    dv_list(ws, rng(C['销售']), '=人员列表', prompt='从【基础资料】的人员里选；提成按人员的比例算')
    dv_list(ws, rng(C['收款账户']), '=账户列表', prompt='定金进哪个账户；空着＝基础资料第一个账户')
    for k in ('订单总金额', '预计成本', '定金金额'):
        dv_dec(ws, rng(C[k]))
    dv_dec(ws, rng(C['提成比例']), '提成比例', '0', '100')
    dv = DataValidation(type='whole', operator='between', formula1='0', formula2='9999', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng(C['人数']))
    dv_len(ws, rng(C['订单号']), 20)
    _example_grey(ws, rows, ORD_R0, C['备注'])
    home_link(ws, f'{CL(CI(ORD_LAST) + 1)}1')
    print_setup(ws, f'{ORD_HDR}:{ORD_HDR}')


# ───────────────────────── 收支登记 ─────────────────────────
def build_cash(ws, ctx):
    C, S = CASH_COLS, CASH_SHOW
    heads = [(C['日期'], '日期'), (C['类别'], '收支类别'), (C['订单号'], '订单号\n（订单的成本、退款要选）'),
             (C['往来对象'], '往来对象\n（供应商 / 员工）'), (C['摘要'], '摘要'), (C['收入'], '收入金额'), (C['支出'], '支出金额'),
             (C['账户'], '账户\n（空＝默认）'), (C['付款情况'], '付款情况\n（空＝已付）'), (C['备注'], '备注'),
             (S['归类'], '归类\n（自动）'), (S['订单客户'], '订单客户\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 定金以外的钱都记这里，一笔一行。订单的钱——收尾款/分期款、退客户款、地接机票酒店等成本、开票税费、收款手续费、渠道返佣——'
           '「订单号」一定要选（下拉只列没结算的单，新的在上面，显示「订单号｜客户｜出行日」；也可以直接打订单号）。'
           '工资、销售提成「往来对象」选人员（工资和提成一起转的也拆成两行）。成本先欠着没付的（比如地接回来再结），付款情况选「未付」：'
           '算进订单成本，不算钱出去；付了以后把「未付」清掉、日期改成付款那天（只付了一部分：原行改成还欠的数，另起一行记付了的）。'
           '现金存银行、微信提现：记两行「内部转账」（转出账户记支出、转入账户记收入）。灰底的是示例行，正式用时整行删掉。可以插行、删行、排序。')
    W = {C['日期']: 11, C['类别']: 12, C['订单号']: 22, C['往来对象']: 13, C['摘要']: 22, C['收入']: 12, C['支出']: 12,
         C['账户']: 10, C['付款情况']: 9, C['备注']: 16, S['归类']: 9, S['订单客户']: 10, S['校验']: 34}
    fm = {C['日期']: DATE, C['收入']: MONEY, C['支出']: MONEY, C['订单号']: '@'}
    al = {C['订单号']: AL, C['往来对象']: AL, C['摘要']: AL, C['备注']: AL, C['收入']: AR, C['支出']: AR, S['校验']: AL}
    rows = ctx.get('cash', [])
    nshow = N_CASH
    _frame(ws, CASH_LAST, CASH_HDR, heads, C_IN, tip, W, CASH_R0, nshow, fm, al, f'C{CASH_R0}', auto=tuple(S.values()))
    ws.row_dimensions[CASH_HDR].height = 42
    for col in S.values():
        ws[f'{col}{CASH_HDR}'].fill = fill('FF8EA9DB')
    last = CASH_R0 + N_CASH + 2000
    put(ws, 'A3', '筛选后看得见的行合计', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells('A3:E3')
    for col in (C['收入'], C['支出']):
        put(ws, f'{col}3', f'=SUBTOTAL(9,{col}{CASH_HDR}:{col}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["账户"]}3', '欠人家没付的', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells(f'{C["账户"]}3:{C["付款情况"]}3')
    put(ws, f'{C["备注"]}3', '=SUMIFS(收_未付额,收_有效,1)', F_AUTOB, FILL_AUTO, MONEY, AR)
    for r in range(CASH_R0 + nshow, CASH_R0 + N_CASH):          # 订单号整列文本格式（容量以内）
        ws[f'{C["订单号"]}{r}'].number_format = '@'
    _write_rows(ws, rows, C, CASH_R0)
    ok = 'AND(INDEX(收_类别,{n})="",INDEX(收_有效,{n})<>1)'
    _show(ws, CASH_R0, nshow, CASH_HDR, {
        S['归类']: f'IF({ok},"",INDEX(收_归类,{{n}}))',
        S['订单客户']: 'INDEX(收_订单客户,{n})&""',
        S['校验']: 'INDEX(收_校验,{n})&""',
    }, {})
    lastshow = CASH_R0 + nshow - 1
    ws.conditional_formatting.add(f'{S["校验"]}{CASH_R0}:{S["校验"]}{lastshow}',
                                  FormulaRule(formula=[f'LEFT({S["校验"]}{CASH_R0},1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(f'A{CASH_R0}:{C["备注"]}{lastshow}',
                                  FormulaRule(formula=[f'${C["付款情况"]}{CASH_R0}="未付"'], font=Font(name=YH, sz=10, color='FFC65911')))
    rng = lambda col: f'{col}{CASH_R0}:{col}{CASH_R0 + N_CASH - 1}'
    dv_date(ws, rng(C['日期']))
    dv_list(ws, rng(C['类别']), '=收支类别列表', prompt='从【基础资料】的收支类别里选；订单的成本选地接费、机票、酒店……')
    dv_list(ws, rng(C['订单号']), '=订单列表', stop=False, prompt='订单成本、分期追加收款、退款给客户要选订单；也可以直接打订单号')
    dv_list(ws, rng(C['往来对象']), '=人员列表', stop=False, prompt='工资、提成选人员；供应商直接打名字')
    dv_list(ws, rng(C['账户']), '=账户列表', prompt='空着＝基础资料第一个账户')
    dv_list(ws, rng(C['付款情况']), '"' + ','.join(PAY_STATUS) + '"', prompt='先欠着没付选「未付」；付了清掉、日期改成付款那天')
    for k in ('收入', '支出'):
        dv_dec(ws, rng(C[k]))
    _example_grey(ws, rows, CASH_R0, C['备注'])
    home_link(ws, f'{CL(CI(CASH_LAST) + 1)}1')
    print_setup(ws, f'{CASH_HDR}:{CASH_HDR}')


# ───────────────────────── 基础资料 ─────────────────────────
def build_base(ws, ctx):
    widths(ws, {'A': 16, 'B': 12, 'C': 12, 'D': 12, 'E': 34, 'F': 3, 'G': 62})
    title(ws, '基 础 资 料', 'G', C_MASTER,
          '💡 一次填好、偶尔改：① 公司名称、建账日期、提醒天数；② 账户（对公、微信、支付宝、现金，填建账日的余额）；'
          '③ 收支类别（收支登记的下拉，「归类」决定算到哪：订单成本、工资、日常费用……）；④ 人员（销售名字、员工，提成按这里的比例算）。'
          '每块上下排着，在块里插整行、删整行都只动这一块；名称不要重复。改了名字，记得用「查找替换」把两张登记表里的旧名字也改掉。')
    home_link(ws, 'G3')
    pv = ctx.get('params', {})
    put(ws, 'A3', '① 参数', F_SEC, fill(C_MASTER), align=AL)
    ws.merge_cells('A3:C3')
    for k, lbl, default in PARAMS:
        r = PAR_ROW[k]
        put(ws, f'A{r}', lbl, F_KPI_L, F_LBL, align=AL)
        ws.merge_cells(f'A{r}:B{r}')
        v = pv.get(k) or default
        if k == '建账日':
            v = _d(v)
        put(ws, f'C{r}', v, F_IN, FILL_IN, DATE if k == '建账日' else None, AC)
    dv_date(ws, f'C{PAR_ROW["建账日"]}')
    for k in ('尾款天数', '结算天数'):
        dv = DataValidation(type='whole', operator='between', formula1='1', formula2='365', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f'C{PAR_ROW[k]}')
    put(ws, f'E{PAR_ROW["公司名称"]}', '首页、订单详情（结算单）抬头用', F_NOTE, align=AL, border=False)
    put(ws, f'E{PAR_ROW["建账日"]}', '从哪天开始用这本账；账户期初余额都是这一天的（建账以前收付的钱不再进资金）', F_NOTE, align=AL, border=False)
    put(ws, f'E{PAR_ROW["尾款天数"]}', '出发前这么多天还没收齐的，首页、未出行订单标红提醒', F_NOTE, align=AL, border=False)
    put(ws, f'E{PAR_ROW["结算天数"]}', '出行后这么多天还没填结算日期的、成本还欠着没付的，提醒', F_NOTE, align=AL, border=False)
    notes = {'账户': '收钱、付钱的地方：对公账户、微信、支付宝、现金……期初余额填建账日的余额。订单的收款账户、收支登记的账户空着＝第一个账户。',
             '类别': '归类决定这类钱算到哪：订单收款（尾款、分期）、订单退款、订单成本（要选订单号）；工资、销售提成（往来对象选人员）；'
                     '日常费用（房租、推广……）；其他收入、其他支出；内部转账（账户之间倒钱）、资金往来（股东投入取出、借还款）不算收支。',
             '人员': '销售名字、员工从这里下拉。提成基数：利润＝按订单利润提（结算了按实际利润，没结算先按预计利润估）；订单金额＝按订单总金额提。'
                     '提成比例填 10% 或 10 都行（1 及以上当百分数）；订单登记里单独填了比例的按订单的。提成算在订单结算的那个月。'}
    heads = {'账户': (ACC_COLS, [('名称', '账户名称'), ('期初余额', '期初余额\n（建账日）'), ('备注', '备注')], ctx.get('accounts', [])),
             '类别': (CAT_COLS, [('名称', '收支类别'), ('方向', '方向'), ('归类', '归类'), ('说明', '说明')], ctx.get('cats', [])),
             '人员': (PER_COLS, [('姓名', '姓名'), ('岗位', '岗位'), ('提成基数', '提成基数'), ('提成比例', '提成比例'), ('备注', '备注')],
                      ctx.get('persons', []))}
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        cols, hd, data = heads[key]
        c2 = cols[hd[-1][0]]
        put(ws, f'A{tr}', ttl, F_SEC, fill(C_MASTER), align=AL)
        ws.merge_cells(f'A{tr}:{c2}{tr}')
        put(ws, f'G{tr}', notes[key], F_NOTE, align=ALW, border=False)
        ws.merge_cells(f'G{tr}:G{hr}')
        header(ws, hr, [(cols[k], t) for k, t in hd], C_MASTER, height=34)
        fm = {}
        if '期初余额' in cols:
            fm[cols['期初余额']] = MONEY
        if '提成比例' in cols:
            fm[cols['提成比例']] = PCT1
        r0 = hr + 1
        style_rows(ws, r0, r0 + cap - 1, [cols[k] for k, _ in hd], fmts=fm,
                   aligns={cols[k]: AL for k in ('名称', '备注', '姓名', '说明') if k in cols})
        for i, x in enumerate(data):
            for k, _t in hd:
                v = x.get(k)
                if v not in (None, ''):
                    ws[f'{cols[k]}{r0 + i}'] = v
        rngB = lambda k: f'{cols[k]}{r0}:{cols[k]}{r0 + cap - 1}'
        if key == '账户':
            dv_dec(ws, rngB('期初余额'))
        if key == '类别':
            dv_list(ws, rngB('方向'), '"' + ','.join(CAT_DIRS) + '"')
            dv_list(ws, rngB('归类'), '"' + ','.join(CAT_KINDS) + '"', prompt='决定这类钱算到哪（见右边说明）')
        if key == '人员':
            dv_list(ws, rngB('岗位'), '"' + ','.join(PER_ROLES) + '"', stop=False)
            dv_list(ws, rngB('提成基数'), '"' + ','.join(COMM_BASES) + '"', prompt='空着＝按利润')
            dv_dec(ws, rngB('提成比例'), '提成比例', '0', '100')
        dv_len(ws, f'A{r0}:A{r0 + cap - 1}', 30)
        put(ws, f'{c2 if key != "类别" else "E"}{r0 + cap}', f'（这一块最多 {cap} 个，超过了【数据校验】会提醒）', F_NOTE, align=AL, border=False)
    ws.freeze_panes = 'A4'
    print_setup(ws, None, landscape=False)


BUILDERS = {SH_ORD: build_ord, SH_CASH: build_cash, SH_BASE: build_base}


def build(wb, ctx):
    for name, fn in BUILDERS.items():
        fn(wb[name], ctx)
