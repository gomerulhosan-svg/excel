# -*- coding: utf-8 -*-
"""查看表（绿）【未出行订单】没有选择格，截止日＝首页截止日 P_截止。
   顶上（第 3～7 行，固定）：大字「全部预计订单未出行预估总利润」＝SUMIFS(单_预计利润, 单_有效,1, 单_出行码,0, 单_有预计,1)；
      旁边：未出行单数、订单总金额、已收、还没收、预计提成；灰字：另有 N 单（订单金额 x）没填预计成本，未计入；
      再一行：已出行没结算 N 单（预计利润、暂算实际利润）；快出发（P_尾款天数 天内）没收齐 N 单 x 元。
   下面是一块「活动区」（第 9 行起），三个区块一个接一个往下排（前一块有几行就占几行）：
   ① 未出行订单（单_出行码＝0，按出行日期排，没填的在最后，最多 600 单）：表头 → 合计 → 明细 → 共 N 单；快出发没收齐的行红色。
   ② 已出行还没结算的（出行码＝1）＋ 已结算还没收齐的（出行码＝2 且还没收＞0），按出行日期排，最多 300 单；
      出行后天数＞P_结算天数 的行红色；暂算实际利润（没结算的）灰色斜体。
   ③ 欠人家的钱（收_未付额≠0，按日期排，最多 300 笔）：小计（欠供应商、应退客户、其他未付、合计）→ 明细；
      出行超过 P_结算天数 天还欠供应商的行红色。右边小表「按往来对象」：未付行里不重复的往来对象（按登记顺序，最多 30 个）。
   隐藏列：AC 第 k 行（常数）、AD 这一行是什么、AE 取第几条、AF 小表第几行、AG 出行码、AH 要不要标红；
      AJ～AN 小表每个往来对象的数；排序键折成 500 行一列的方块（SMALL 对整块取第 k 小，MOD(…,10000) 还原第几条）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'P'
R0 = 9
NL1, NL2, NL3, NL4 = 600, 300, 300, 30
NREG = 3 + NL1 + 2 + 3 + NL2 + 2 + 6 + NL3 + 1
R1 = R0 + NREG - 1
BLK = 500
SL, SV = 'AA', 'AB'                                   # 标量：说明、值（第 3 行起）
HK, HC, HI, HT, HD, HR = 'AC', 'AD', 'AE', 'AF', 'AG', 'AH'
ENT = dict(名='AJ', 欠='AK', 退='AL', 其他='AM', 合计='AN')   # 小表：第 t 个往来对象在第 R0+t-1 行
K1C, K2C, K3C, K4C = 'AP', 'AW', 'BD', 'BU'          # 排序键方块起始列

SC = ['c1', 'n1', 'c2', 'n2', 'c3', 'n3', 'e4', 'm4', 'b2', 'b3', '最后k', '最后一行',
      '预估总利润', '未出行单数', '未出行金额', '未出行已收', '未出行还没收', '未出行预计成本', '未出行已记成本', '未出行预计提成',
      '没预计单数', '没预计金额', '待结单数', '待结预计利润', '待结暂算利润', '快出发单数', '快出发还没收',
      '②金额', '②已收', '②还没收', '②已记成本', '②未付成本', '②预计利润', '②实际利润',
      '欠供应商', '应退客户', '其他未付', '未付合计', '表欠余', '表退余', '表其他余', '表合计余', 'o4', 'tN']
SR = {k: 3 + i for i, k in enumerate(SC)}

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
GREEN_H = 'FF70AD47'
F_WHITE_B = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_GREY = Font(name=YH, sz=9, color='FF808080')
F_GREY_I = Font(name=YH, sz=10, italic=True, color='FF9E9E9E')
F_OVER = Font(name=YH, sz=10, bold=True, color='FF9C0006')
F_BIG = Font(name=YH, sz=20, bold=True, color='FF1F3864')


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def cf_blocks(ws, r0, r1, rules):
    """条件格式：rules＝[(起列, 止列, cond(块首列)→公式, 样式)]，按优先顺序。
       把列切成「规则集合一样」的连续块，每格只落在一个块里、块里规则按顺序排（都 stopIfTrue）：
       Excel 按优先级取第一条成立的；LibreOffice / 有的 WPS 版本对重叠的几个块只认一个，这样写哪边都一样。"""
    lo = min(CI(a) for a, _b, _c, _k in rules)
    hi = max(CI(b) for _a, b, _c, _k in rules)
    runs = []
    for i in range(lo, hi + 1):
        sig = tuple(j for j, (a, b, _c, _k) in enumerate(rules) if CI(a) <= i <= CI(b))
        if runs and runs[-1][2] == sig:
            runs[-1][1] = i
        else:
            runs.append([i, i, sig])
    for i0, i1, sig in runs:
        for j in sig:
            _a, _b, cond, kw = rules[j]
            ws.conditional_formatting.add(f'{CL(i0)}{r0}:{CL(i1)}{r1}',
                                          FormulaRule(formula=[cond(CL(i0))], stopIfTrue=True, **kw))


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def _q(t):
    return '"' + t.replace('"', '""') + '"'


def key_block(ws, c0, n, cell):
    """排序键折成 BLK 行一列的方块：第 i 条（0 起）在 c0 右边第 i//BLK 列、第 R0+i%BLK 行；cell(n) 给出第 n 条的公式（不含 =）。
       返回方块区域（绝对引用）"""
    i0 = CI(c0)
    for i in range(n):
        c = ws[f'{CL(i0 + i // BLK)}{R0 + i % BLK}']
        c.value = '=' + cell(i + 1)
        c.font = F_HELP
    ncol = -(-n // BLK)
    return f'${c0}${R0}:${CL(i0 + ncol - 1)}${R0 + BLK - 1}'


TIP = ('💡 全自动不用填，截止日期＝【首页】的截止日期。最上面大字是用户要的「全部预计订单未出行预估总利润」＝还没出行（没取消）的订单的预计利润'
       '（订单总金额－预计成本）合计；没填预计成本的不算，另外列出来。下面三块一块接一块往下排：'
       '① 未出行订单，按出行日期排（没填出行日期的在最后），红色＝出发前 N 天内（【基础资料】里设，默认 15 天）还没收齐；'
       '② 已经出行还没结算的、已经结算但钱还没收齐的，红色＝出行超过 N 天（默认 30 天）——该收尾款、对成本、填结算日期了；'
       '暂算实际利润（灰色斜体）＝订单总金额－这单已记的成本（含未付的），结算后才定下来；'
       '③ 欠人家的钱：【收支登记】里付款情况选了「未付」的每一笔（欠地接、机票等供应商的，答应退客户还没退的，别的没付的），'
       '红色＝出行超过 N 天还欠供应商；右边按往来对象合计。付了以后去【收支登记】把「未付」清掉、日期改成付款那天，这里就没了。')


def build(wb, ctx=None):
    ws = wb[SH_TODO]
    W = {'A': 5, 'B': 12, 'C': 9, 'D': 10, 'E': 12, 'F': 16, 'G': 11, 'H': 12, 'I': 13, 'J': 12, 'K': 12, 'L': 12,
         'M': 12, 'N': 12, 'O': 12, 'P': 14}
    widths(ws, W)
    title(ws, '未出行订单（预估总利润 · 没结算的 · 欠人家的钱）', LAST, C_VIEW, TIP)

    # ── 排序键（方块） ──
    u_ = lambda k, n: f'INDEX(单_{k},{n})'
    s_ = lambda k, n: f'INDEX(收_{k},{n})'
    K1 = key_block(ws, K1C, N_ORD, lambda n: (f'IF(AND({u_("有效", n)}=1,{u_("出行码", n)}=0),{u_("出行键", n)},"")'))
    K2 = key_block(ws, K2C, N_ORD, lambda n: (f'IF({u_("有效", n)}<>1,"",IF(OR({u_("出行码", n)}=1,'
                                              f'AND({u_("出行码", n)}=2,{u_("还没收", n)}>0)),{u_("出行键", n)},""))'))
    K3 = key_block(ws, K3C, N_CASH, lambda n: f'IF({s_("未付额", n)}<>0,{s_("排序键", n)},"")')
    K4 = key_block(ws, K4C, N_CASH, lambda n: (f'IF({s_("未付额", n)}=0,"",IF({s_("往来对象", n)}="","",'
                                               f'IF(COUNTIFS(收_往来对象,{esc(s_("往来对象", n))},收_未付额,"<>0",收_n,"<"&{n})=0,{n},"")))'))
    for c, t in ((K1C, '①键'), (K2C, '②键'), (K3C, '③键'), (K4C, '往来对象第一次'), (HK, 'k'), (HC, '这行是'), (HI, '第几条'),
                 (HT, '小表行'), (HD, '出行码'), (HR, '标红'), (ENT['名'], '往来对象'), (ENT['欠'], '欠供应商'), (ENT['退'], '应退客户'),
                 (ENT['其他'], '其他未付'), (ENT['合计'], '合计')):
        ws[f'{c}{R0 - 1}'] = t
        ws[f'{c}{R0 - 1}'].font = F_HELP

    # ── 小表：每个往来对象（第 t 个＝未付行里第 t 个第一次出现的） ──
    for t in range(1, NL4 + 1):
        r = R0 + t - 1
        nm = f'${ENT["名"]}{r}'
        ws[f'{ENT["名"]}{r}'] = f'=IF({t}>{S("e4")},"",INDEX(收_往来对象,SMALL({K4},{t})))'
        base = f'收_未付额,收_有效,1,收_往来对象,{esc(nm)}'
        ws[f'{ENT["欠"]}{r}'] = f'=IF({nm}="",0,ROUND(SUMIFS({base},收_归类,"订单成本"),2))'
        ws[f'{ENT["退"]}{r}'] = f'=IF({nm}="",0,ROUND(SUMIFS({base},收_归类,"订单退款"),2))'
        ws[f'{ENT["合计"]}{r}'] = f'=IF({nm}="",0,ROUND(SUMIFS({base}),2))'
        ws[f'{ENT["其他"]}{r}'] = f'=ROUND(${ENT["合计"]}{r}-${ENT["欠"]}{r}-${ENT["退"]}{r},2)'
        for c in ENT.values():
            ws[f'{c}{r}'].font = F_HELP
    ent = lambda key: f'${ENT[key]}${R0}:${ENT[key]}${R0 + NL4 - 1}'

    # ── 标量 ──
    o1 = '单_有效,1,单_出行码,0'
    two = lambda f: (f'ROUND(SUMIFS({f},单_有效,1,单_出行码,1)+SUMIFS({f},单_有效,1,单_出行码,2,单_还没收,">0"),2)')
    soon = f'{o1},单_出行日期,">0",单_出行日期,"<="&(P_截止+P_尾款天数),单_还没收,">0"'
    sc = {
        'c1': f'=COUNT({K1})', 'n1': f'=MIN({S("c1")},{NL1})',
        'c2': f'=COUNT({K2})', 'n2': f'=MIN({S("c2")},{NL2})',
        'c3': f'=COUNT({K3})', 'n3': f'=MIN({S("c3")},{NL3})',
        'e4': f'=COUNT({K4})', 'm4': f'=MIN({S("e4")},{NL4})',
        'b2': f'=5+{S("n1")}', 'b3': f'={S("b2")}+5+{S("n2")}', '最后k': f'={S("b3")}+7+{S("n3")}',
        '最后一行': f'={R0 - 1}+{S("最后k")}',
        '预估总利润': f'=ROUND(SUMIFS(单_预计利润,{o1},单_有预计,1),2)',
        '未出行单数': f'=COUNTIFS({o1})',
        '未出行金额': f'=ROUND(SUMIFS(单_订单总金额,{o1}),2)',
        '未出行已收': f'=ROUND(SUMIFS(单_已收,{o1}),2)',
        '未出行还没收': f'=ROUND(SUMIFS(单_还没收,{o1}),2)',
        '未出行预计成本': f'=ROUND(SUMIFS(单_预计成本,{o1},单_有预计,1),2)',
        '未出行已记成本': f'=ROUND(SUMIFS(单_实际成本,{o1}),2)',
        '未出行预计提成': f'=ROUND(SUMIFS(单_订单提成,{o1}),2)',
        '没预计单数': f'=COUNTIFS({o1},单_有预计,0)',
        '没预计金额': f'=ROUND(SUMIFS(单_订单总金额,{o1},单_有预计,0),2)',
        '待结单数': '=COUNTIFS(单_有效,1,单_出行码,1)',
        '待结预计利润': '=ROUND(SUMIFS(单_预计利润,单_有效,1,单_出行码,1),2)',
        '待结暂算利润': '=ROUND(SUMIFS(单_实际利润,单_有效,1,单_出行码,1),2)',
        '快出发单数': f'=COUNTIFS({soon})',
        '快出发还没收': f'=ROUND(SUMIFS(单_还没收,{soon}),2)',
        '②金额': '=' + two('单_订单总金额'), '②已收': '=' + two('单_已收'), '②还没收': '=' + two('单_还没收'),
        '②已记成本': '=' + two('单_实际成本'), '②未付成本': '=' + two('单_未付成本'),
        '②预计利润': '=' + two('单_预计利润'), '②实际利润': '=' + two('单_实际利润'),
        '欠供应商': '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单成本"),2)',
        '应退客户': '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单退款"),2)',
        '未付合计': '=ROUND(SUMIFS(收_未付额,收_有效,1),2)',
        '其他未付': f'=ROUND({S("未付合计")}-{S("欠供应商")}-{S("应退客户")},2)',
        '表欠余': f'=ROUND({S("欠供应商")}-SUM({ent("欠")}),2)',
        '表退余': f'=ROUND({S("应退客户")}-SUM({ent("退")}),2)',
        '表其他余': f'=ROUND({S("其他未付")}-SUM({ent("其他")}),2)',
        '表合计余': f'=ROUND({S("未付合计")}-SUM({ent("合计")}),2)',
        'o4': f'=IF(OR({S("表欠余")}<>0,{S("表退余")}<>0,{S("表其他余")}<>0,{S("表合计余")}<>0),1,0)',
        'tN': f'={S("m4")}+{S("o4")}+1',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    # ── 顶上（第 3～7 行） ──
    lbl_fill = fill('FFD9E1F2')
    put(ws, 'A3', '截止日期', F_KPI_L, lbl_fill, align=AC)
    ws.merge_cells('A3:B3')
    put(ws, 'B3', None, F_KPI_L, lbl_fill)
    put(ws, 'C3', '=P_截止', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('C3:D3')
    put(ws, 'D3', None, F_AUTOB, FILL_AUTO)
    c = put(ws, 'E3', '要换日期去【首页】改 →', Font(name=YH, sz=9, color='FF0563C1', underline='single'), align=AL, border=False)
    link(c, SH_HOME, HOME_END)
    ws.merge_cells('E3:G3')
    ws.merge_cells('H3:O3')
    put(ws, 'H3', '="提醒天数（【基础资料】里改）：出发前 "&P_尾款天数&" 天还没收齐标红；出行超过 "&P_结算天数&" 天还没结算、还欠供应商标红"',
        F_NOTE, align=AL, border=False)
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 22

    ws.merge_cells('A4:F5')
    put(ws, 'A4', '全部预计订单\n未出行预估总利润', Font(name=YH, sz=13, bold=True, color='FF1F3864'), lbl_fill, align=ACW)
    ws.merge_cells('G4:I5')
    put(ws, 'G4', f'={S("预估总利润")}', F_BIG, FILL_TOT, MONEY, AC)
    for rr in (4, 5):
        for col in 'ABCDEFGHI':
            ws[f'{col}{rr}'].border = BD
    kp = [('J', '未出行单数', S('未出行单数'), '0'), ('K', '订单总金额', S('未出行金额'), MONEY), ('L', '已收', S('未出行已收'), MONEY),
          ('M', '还没收', S('未出行还没收'), MONEY), ('N', '预计提成', S('未出行预计提成'), MONEY)]
    for col, lbl, f, fmt in kp:
        put(ws, f'{col}4', lbl, F_KPI_L, lbl_fill, align=ACW)
        put(ws, f'{col}5', f'={f}', Font(name=YH, sz=13, bold=True, color='FF1F3864'), FILL_AUTO, fmt, AC)
    ws.merge_cells('O4:P5')
    put(ws, 'O4', '预估总利润＝未出行（没取消）订单的「订单总金额－预计成本」合计', F_NOTE, align=ALW)
    for rr in (4, 5):
        for col in 'OP':
            ws[f'{col}{rr}'].border = BD
    ws.row_dimensions[4].height = 24
    ws.row_dimensions[5].height = 30
    ws.merge_cells(f'A6:{LAST}6')
    put(ws, 'A6', (f'=IF({S("没预计单数")}=0,"未出行的订单都填了预计成本 ✓",'
                   f'"另有 "&{S("没预计单数")}&" 单（订单金额 "&TEXT({S("没预计金额")},"#,##0.00")&"）没填预计成本，未计入上面的预估总利润'
                   f'——去【订单登记】补上预计成本")'), F_NOTE, align=AL, border=False)
    ws.merge_cells(f'A7:{LAST}7')
    put(ws, 'A7', (f'="已出行没结算 "&{S("待结单数")}&" 单，预计利润 "&TEXT({S("待结预计利润")},"#,##0.00")'
                   f'&"、暂算实际利润 "&TEXT({S("待结暂算利润")},"#,##0.00")'
                   f'&"；快出发（"&P_尾款天数&" 天内）没收齐 "&{S("快出发单数")}&" 单 "&TEXT({S("快出发还没收")},"#,##0.00")&" 元"'),
        F_TXTB, align=AL, border=False)
    ws.conditional_formatting.add('A7', FormulaRule(formula=[f'OR({S("快出发单数")}>0,{S("待结单数")}>0)'],
                                                    font=Font(bold=True, color='FFC00000')))
    ws.conditional_formatting.add('A6', FormulaRule(formula=[f'{S("没预计单数")}>0'], font=Font(color='FFC65911')))
    ws.row_dimensions[6].height = 18
    ws.row_dimensions[7].height = 20
    ws.row_dimensions[8].height = 8

    # ── 活动区 ──
    n1, n2, n3 = S('n1'), S('n2'), S('n3')
    b2, b3 = S('b2'), S('b3')
    fm = {'A': '0', 'B': DATE, 'C': '0', 'G': INT, 'I': MONEY, 'J': MONEY, 'K': MONEY, 'L': MONEY, 'M': MONEY,
          'N': MONEY, 'O': MONEY, 'P': MONEY}
    al = {'A': AC, 'B': AC, 'C': AC, 'D': AC, 'E': AL, 'F': AL, 'G': AC, 'H': AL, 'I': AR, 'J': AR, 'K': AR, 'L': AR,
          'M': AR, 'N': AR, 'O': AR, 'P': AC}
    order_kinds = ('D1', 'D2', 'D3', 'S1', 'S2', 'S3', 'H1', 'H2', 'H3', 'T1', 'T2', 'Ta', 'Tb', 'Tc', 'Td', 'N1', 'N2', 'N3')
    for kk in range(NREG):
        r = R0 + kk
        k = f'${HK}{r}'
        cd, ix, tj = f'${HC}{r}', f'${HI}{r}', f'${HT}{r}'
        ws[f'{HK}{r}'] = kk + 1
        ws[f'{HC}{r}'] = (f'=IF({k}<=3+{n1},CHOOSE(MIN({k},4),"S1","H1","T1","D1"),IF({k}=4+{n1},"N1",IF({k}<={b2},"",'
                          f'IF({k}<={b2}+3,CHOOSE({k}-{b2},"S2","H2","T2"),IF({k}<={b2}+3+{n2},"D2",IF({k}={b2}+4+{n2},"N2",'
                          f'IF({k}<={b3},"",IF({k}<={b3}+6,CHOOSE({k}-{b3},"S3","H3","Ta","Tb","Tc","Td"),'
                          f'IF({k}<={b3}+6+{n3},"D3",IF({k}={b3}+7+{n3},"N3",""))))))))))')
        ws[f'{HI}{r}'] = (f'=IF({cd}="D1",MOD(SMALL({K1},{k}-3),10000),IF({cd}="D2",MOD(SMALL({K2},{k}-{b2}-3),10000),'
                          f'IF({cd}="D3",MOD(SMALL({K3},{k}-{b3}-6),10000),0)))')
        ws[f'{HT}{r}'] = f'=IF(AND({k}>{b3}+2,{k}<={b3}+2+{S("tN")}),{k}-{b3}-2,0)'
        u = lambda key: f'INDEX(单_{key},{ix})'
        s = lambda key: f'INDEX(收_{key},{ix})'
        orow = s('订单行')
        ws[f'{HD}{r}'] = f'=IF({cd}="D2",{u("出行码")},"")'
        ws[f'{HR}{r}'] = '=' + _chain([
            (f'{cd}="D1"', f'IF(AND({u("出行日期")}>0,{u("出行日期")}-P_截止<=P_尾款天数,{u("还没收")}>0),1,0)'),
            (f'{cd}="D2"', f'IF(N({u("出行天数")})>P_结算天数,1,0)'),
            (f'{cd}="D3"', (f'IF(AND({s("归类")}="订单成本",{orow}>0),IF(AND(INDEX(单_出行码,{orow})<>3,'
                            f'N(INDEX(单_出行天数,{orow}))>P_结算天数),1,0),0)')),
        ], '0')
        for c in (HK, HC, HI, HT, HD, HR):
            ws[f'{c}{r}'].font = F_HELP

        tdate = f'IF({u("出行日期")}>0,{u("出行日期")},"没填")'
        est = lambda key: f'IF({u("有预计")}=1,{u(key)},"没填")'
        kind3 = f'IF({s("归类")}="订单成本","欠供应商",IF({s("归类")}="订单退款","应退客户","其他未付"))'
        cust3 = (f'IF({orow}>0,INDEX(单_客户名字,{orow})&IF(INDEX(单_出行日期,{orow})>0," · 出行 "'
                 f'&TEXT(INDEX(单_出行日期,{orow}),"m/d"),""),"")')
        stat2 = (f'IF({u("出行码")}=1,IF(N({u("出行天数")})>P_结算天数,"超"&P_结算天数&"天没结算","已出行没结算"),"已结算没收齐")')
        nrow = lambda c, unit, cap: f'IF({S(c)}=0,"（没有）","共 "&{S(c)}&" {unit}")'
        capw = lambda c, unit, cap: f'IF({S(c)}>{cap},"⚠ 只列前 {cap} {unit}","")'
        D = {
            'A': {'D1': f'{k}-3', 'D2': f'{k}-{b2}-3', 'D3': f'{k}-{b3}-6', 'S1': '"①"', 'S2': '"②"', 'S3': '"③"',
                  'H1': '"序号"', 'H2': '"序号"', 'H3': '"序号"'},
            'B': {'D1': tdate, 'D2': tdate, 'D3': s('日期'), 'S1': '"未出行订单"', 'S2': '"已出行没结清"', 'S3': '"欠人家的钱"',
                  'H1': '"出行日期"', 'H2': '"出行日期"', 'H3': '"日期"', 'T1': '"合计"', 'T2': '"合计"', 'Ta': '"小计"',
                  'N1': nrow('c1', '单', NL1), 'N2': nrow('c2', '单', NL2), 'N3': nrow('c3', '笔', NL3)},
            'C': {'D1': f'IF({u("出行日期")}>0,{u("出行日期")}-P_截止,"")', 'D2': u('出行天数'), 'D3': kind3,
                  'H1': '"还有几天"', 'H2': '"出行后天数"', 'H3': '"分类"',
                  'Ta': '"欠供应商"', 'Tb': '"应退客户"', 'Tc': '"其他未付"', 'Td': '"合计"'},
            'D': {'D1': u('订单号'), 'D2': u('订单号'), 'D3': s('类别'), 'H1': '"订单号"', 'H2': '"订单号"', 'H3': '"收支类别"'},
            'E': {'D1': u('客户名字'), 'D2': u('客户名字'), 'D3': s('订单号'), 'H1': '"客户名字"', 'H2': '"客户名字"', 'H3': '"订单号"'},
            'F': {'D1': u('线路'), 'D2': u('线路'), 'D3': cust3, 'S1': '"按出行日期排"', 'S2': '"按出行日期排"', 'S3': '"付款情况＝未付"',
                  'H1': '"线路 / 目的地"', 'H2': '"线路 / 目的地"', 'H3': '"订单客户"',
                  'N1': capw('c1', '单', NL1), 'N2': capw('c2', '单', NL2), 'N3': capw('c3', '笔', NL3)},
            'G': {'D1': u('人数'), 'D2': u('人数'), 'D3': s('往来对象'), 'H1': '"人数"', 'H2': '"人数"', 'H3': '"往来对象"'},
            'H': {'D1': u('销售'), 'D2': u('销售'), 'D3': s('摘要'), 'H1': '"销售名字"', 'H2': '"销售名字"', 'H3': '"摘要"'},
            'I': {'D1': u('订单总金额'), 'D2': u('订单总金额'), 'D3': s('未付额'), 'H1': '"订单总金额"', 'H2': '"订单总金额"',
                  'H3': '"金额"', 'T1': S('未出行金额'), 'T2': S('②金额'), 'Ta': S('欠供应商'), 'Tb': S('应退客户'),
                  'Tc': S('其他未付'), 'Td': S('未付合计')},
            'J': {'D1': u('已收'), 'D2': u('已收'), 'D3': f'"第 "&{s("录入行")}&" 行"', 'H1': '"已收"', 'H2': '"已收"',
                  'H3': '"收支登记行号"', 'T1': S('未出行已收'), 'T2': S('②已收')},
            'K': {'D1': u('还没收'), 'D2': u('还没收'), 'H1': '"还没收"', 'H2': '"还没收"', 'T1': S('未出行还没收'), 'T2': S('②还没收')},
            'L': {'D1': est('预计成本'), 'D2': u('实际成本'), 'H1': '"预计成本"', 'H2': '"已记成本"', 'H3': '"往来对象"',
                  'S3': '"按往来对象"', 'T1': S('未出行预计成本'), 'T2': S('②已记成本')},
            'M': {'D1': u('实际成本'), 'D2': u('未付成本'), 'H1': '"已记成本"', 'H2': '"未付成本"', 'H3': '"欠供应商"',
                  'T1': S('未出行已记成本'), 'T2': S('②未付成本')},
            'N': {'D1': est('预计利润'), 'D2': est('预计利润'), 'H1': '"预计利润"', 'H2': '"预计利润"', 'H3': '"应退客户"',
                  'T1': S('预估总利润'), 'T2': S('②预计利润')},
            'O': {'D1': u('订单提成'), 'D2': u('实际利润'), 'H1': '"预计提成"', 'H2': '"暂算实际利润"', 'H3': '"其他未付"',
                  'T1': S('未出行预计提成'), 'T2': S('②实际利润')},
            'P': {'D1': u('尾款收否'), 'D2': stat2, 'H1': '"尾款收否"', 'H2': '"状态"', 'H3': '"合计"'},
        }
        m4, o4, tN = S('m4'), S('o4'), S('tN')
        other = f'AND({o4}=1,{tj}={m4}+1)'
        TB = {
            'L': f'IF({tj}<={m4},INDEX({ent("名")},{tj}),IF({other},IF({S("e4")}>{NL4},"其余往来对象","（没填）"),"合计"))',
            'M': f'IF({tj}<={m4},INDEX({ent("欠")},{tj}),IF({other},{S("表欠余")},{S("欠供应商")}))',
            'N': f'IF({tj}<={m4},INDEX({ent("退")},{tj}),IF({other},{S("表退余")},{S("应退客户")}))',
            'O': f'IF({tj}<={m4},INDEX({ent("其他")},{tj}),IF({other},{S("表其他余")},{S("其他未付")}))',
            'P': f'IF({tj}<={m4},INDEX({ent("合计")},{tj}),IF({other},{S("表合计余")},{S("未付合计")}))',
        }
        for col, m in D.items():
            body = _chain([(f'{cd}="{kd}"', m[kd]) for kd in order_kinds if kd in m])
            if col in TB:
                body = f'IF({tj}>0,{TB[col]},{body})'
            x = ws[f'{col}{r}']
            x.value = '=' + body
            x.font = F_TXT
            x.alignment = al[col]
            if col in fm:
                x.number_format = fm[col]

    # ── 条件格式（按隐藏列 AD 这一行是什么） ──
    c0, t0, red0, code0 = f'${HC}{R0}', f'${HT}{R0}', f'${HR}{R0}', f'${HD}{R0}'
    is_ = lambda *ks: 'OR(' + ','.join(f'{c0}="{x}"' for x in ks) + ')'
    A, J, L = 'A', 'J', 'L'
    cf_blocks(ws, R0, R1, [
        (A, LAST, lambda c: is_('S1', 'S2', 'S3'), dict(fill=cf_fill(C_VIEW), font=F_WHITE_B)),
        (A, LAST, lambda c: f'AND({is_("H1", "H2", "H3")},{c}{R0}<>"")', dict(fill=cf_fill(GREEN_H), font=F_WHITE_B, border=CF_BD)),
        (L, LAST, lambda c: f'AND({t0}>0,{t0}={S("tN")})', dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD)),
        (L, LAST, lambda c: f'{t0}>0', dict(border=CF_BD)),
        (A, LAST, lambda c: is_('T1', 'T2'), dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD)),
        (A, J, lambda c: f'{c0}="Td"', dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD)),
        (A, J, lambda c: is_('Ta', 'Tb', 'Tc'), dict(fill=cf_fill('FFDDEBF7'), border=CF_BD)),
        (A, LAST, lambda c: f'AND({red0}=1,{is_("D1", "D2")})', dict(fill=cf_fill('FFFFC7CE'), font=F_OVER, border=CF_BD)),
        (A, J, lambda c: f'AND({red0}=1,{c0}="D3")', dict(fill=cf_fill('FFFFC7CE'), font=F_OVER, border=CF_BD)),
        ('O', 'O', lambda c: f'OR({c0}="D1",AND({c0}="D2",{code0}=1))', dict(font=F_GREY_I, border=CF_BD)),
        (L, 'N', lambda c: f'AND({is_("D1", "D2")},{c}{R0}="没填")', dict(font=F_GREY_I, border=CF_BD)),
        (A, LAST, lambda c: is_('D1', 'D2'), dict(border=CF_BD)),
        (A, J, lambda c: f'{c0}="D3"', dict(border=CF_BD)),
        (A, LAST, lambda c: is_('N1', 'N2', 'N3'), dict(font=F_GREY)),
    ])

    hide(ws, *[CL(i) for i in range(CI(SL), CI(K4C) + N_CASH // BLK)])
    ws.freeze_panes = 'A4'
    print_setup(ws, None, landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("最后一行")})')
