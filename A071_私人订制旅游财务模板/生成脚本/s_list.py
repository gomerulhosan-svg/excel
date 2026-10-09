# -*- coding: utf-8 -*-
"""查看表（绿）【订单清单】用户要的「每个订单详细清单」。
   黄格（第 3 行，第 4 行灰字是实际用的值）：B3 按哪个日期（预订日期/出行日期；空＝预订日期）、D3 起、F3 止（空＝不限）、
   H3 销售（空＝全部）、J3 状态（全部/未出行/已出行/已结算/已取消/没收齐；空＝全部）、L3 客户（包含这几个字；空＝全部）。
   合计区（第 5～7 行，不随清单移动）：单数、订单总金额、定金、已收、还没收、预计利润（只算填了预计成本的）、
   实际利润（只算已出行/已结算/已取消）、订单提成、其中应发。
   清单（第 8 行表头，第 9 行起）：前 12 列按用户原话的顺序；按所选日期从早到晚排（同一天按登记顺序；没填这个日期的排最后）。
   排序键：本表右边隐藏列 AD（每个订单一格，共 N_ORD 格）＝日期（没填＝99999）×10000＋第几条，不符合条件＝空；
   清单第 k 行＝MOD(SMALL(排序键,k),10000)。合计＝SUMIFS(单_xxx, 单_有效,1, 排序键,">0")，跟清单同一个条件。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'S'
HDR, R0 = 8, 9
SHOW = 1500                                   # 清单最多显示几单（下面再多一行「共 N 单」）
NROW = SHOW + 1
SL, SV = 'AA', 'AB'                           # 标量：说明、值（第 3 行起）
IDX = 'AC'                                    # 清单每行：取第几条订单（0＝空行）
KEY = 'AD'                                    # 排序键：第 n 条订单在第 R0+n-1 行
KEY_RNG = f'${KEY}${R0}:${KEY}${R0 + N_ORD - 1}'
STATUS = ['全部', '未出行', '已出行', '已结算', '已取消', '没收齐']
SC = ['日期选', '模式', '起', '止', '有日期条件', '止用', '销售', '状态选', '状态码', '客户', '单数', '显示数',
      '没填预计', '已取消单数', '已取消金额', '暂算利润', '预估提成', '最后一行']
SR = {k: 3 + i for i, k in enumerate(SC)}

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_GREY_I = Font(name=YH, sz=10, italic=True, color='FF9E9E9E')
F_GREY = Font(name=YH, sz=10, color='FF9E9E9E')
F_ORANGE_B = Font(name=YH, sz=10, bold=True, color='FFC65911')
C_HDR2 = 'FF548235'                            # 用户原话那 12 列表头深绿，后面补充的列浅一点
C_HDR3 = 'FF70AD47'


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def cf_blocks(ws, r0, r1, rules):
    """条件格式：rules＝[(起列, 止列, cond(块首列)→公式, 样式)]，按优先顺序。
       把列切成「规则集合一样」的连续块，每格只落在一个块里、块里规则按顺序排（都 stopIfTrue）：
       Excel 按优先级取第一条成立的；LibreOffice / 有的 WPS 版本对重叠的几个块只认一个，这样写哪边都一样。"""
    lo = min(CI(a) for a, _b, _c, _k in rules)
    hi = max(CI(b) for _a, b, _c, _k in rules)
    runs = []
    for i in range(lo, hi + 1):
        sig = tuple(j for j, (a, b, _c, _k) in enumerate(rules) if CI(a) <= i <= CI(b))
        if runs and runs[-1][2] == sig:
            runs[-1][1] = i
        else:
            runs.append([i, i, sig])
    for i0, i1, sig in runs:
        for j in sig:
            _a, _b, cond, kw = rules[j]
            ws.conditional_formatting.add(f'{CL(i0)}{r0}:{CL(i1)}{r1}',
                                          FormulaRule(formula=[cond(CL(i0))], stopIfTrue=True, **kw))


TIP = ('💡 每个订单一行（用户要的「每个订单详细清单」），全自动不用填，要改订单去【订单登记】改。黄格可以筛：按哪个日期（预订日期 / 出行日期）、'
       '起止（空着＝不限）、销售、状态（未出行 / 已出行 / 已结算 / 已取消 / 没收齐＝没取消、还有钱没收）、客户（填名字里的几个字）；黄格都空着＝全部订单。'
       '按选的日期从早到晚排（同一天按登记顺序，没填这个日期的排最后）。上面合计是筛出来的全部订单（不受「只列前 1500 单」限制）。'
       '预计利润「没填」＝订单登记没填预计成本（合计不算）；实际利润：未出行的空着，已出行没结算的是暂算（灰色斜体），结算后才定下来；'
       '订单提成：没结算的是预估（灰色斜体），结算了才算应发。已取消的整行灰色；尾款没收齐的橙色。已收、出行状态都按【首页】的截止日期算。')


def build(wb, ctx=None):
    ws = wb[SH_LIST]
    W = {'A': 8, 'B': 11, 'C': 14, 'D': 11, 'E': 12, 'F': 11, 'G': 13, 'H': 12, 'I': 12, 'J': 11, 'K': 9, 'L': 12,
         'M': 18, 'N': 5, 'O': 12, 'P': 12, 'Q': 12, 'R': 8, 'S': 11}
    widths(ws, W)
    title(ws, '订单清单（每个订单的详细清单）', LAST, C_VIEW, TIP)

    # ── 第 3 行选择格、第 4 行实际用的值 ──
    selector(ws, 'A3', '按哪个\n日期', 'B3', None, '"预订日期,出行日期"', prompt='按预订日期还是出行日期筛、排；空着＝预订日期')
    ws['A3'].alignment = ACW
    selector(ws, 'C3', '起', 'D3', None, fmt=DATE)
    selector(ws, 'E3', '止', 'F3', None, fmt=DATE)
    selector(ws, 'G3', '销售', 'H3', None, '=人员列表', prompt='选一个销售；空着＝全部')
    selector(ws, 'I3', '状态', 'J3', None, '"' + ','.join(STATUS) + '"',
             prompt='未出行 / 已出行（没结算）/ 已结算 / 已取消 / 没收齐（没取消、还有钱没收）；空着＝全部')
    selector(ws, 'K3', '客户', 'L3', None, prompt='填客户名字里的几个字；空着＝全部客户')
    dv_date(ws, 'D3')
    dv_date(ws, 'F3')
    ws.row_dimensions[3].height = 30
    home_link(ws, f'{LAST}3')

    sc = {
        '日期选': '=TRIM(B3&"")',
        '模式': f'=IF({S("日期选")}="出行日期",2,1)',
        '起': '=INT(N(D3))',
        '止': '=INT(N(F3))',
        '有日期条件': f'=IF(OR({S("起")}>0,{S("止")}>0),1,0)',
        '止用': f'=IF({S("止")}>0,{S("止")},99999)',
        '销售': '=TRIM(H3&"")',
        '状态选': '=TRIM(J3&"")',
        '状态码': '=' + ''.join(f'IF({S("状态选")}="{t}",{i},' for i, t in enumerate(STATUS[1:], 1)) + '0' + ')' * (len(STATUS) - 1),
        '客户': '=TRIM(L3&"")',
        '单数': f'=COUNT({KEY_RNG})',
        '显示数': f'=MIN({S("单数")},{SHOW})',
        '没填预计': f'=COUNTIFS(单_有效,1,{KEY_RNG},">0",单_有预计,0)',
        '已取消单数': f'=COUNTIFS(单_有效,1,{KEY_RNG},">0",单_出行码,3)',
        '已取消金额': f'=SUMIFS(单_订单总金额,单_有效,1,{KEY_RNG},">0",单_出行码,3)',
        '暂算利润': f'=ROUND(SUMIFS(单_实际利润,单_有效,1,{KEY_RNG},">0",单_出行码,1),2)',
        '预估提成': f'=ROUND(SUMIFS(单_订单提成,单_有效,1,{KEY_RNG},">0",单_应发,0),2)',
        '最后一行': f'={R0}+{S("显示数")}',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    put(ws, 'A4', '实际用', F_NOTE, align=AC, border=False)
    shows = {
        'B4': (f'=IF({S("模式")}=2,"出行日期","预订日期")', None),
        'D4': (f'=IF({S("起")}>0,{S("起")},"不限")', DATE),
        'F4': (f'=IF({S("止")}>0,{S("止")},"不限")', DATE),
        'H4': (f'=IF({S("销售")}="","全部销售",{S("销售")})', None),
        'J4': (f'=IF({S("状态码")}=0,"全部",{S("状态选")})', None),
        'L4': (f'=IF({S("客户")}="","全部客户","含「"&{S("客户")}&"」")', None),
    }
    for c, (f, fmt) in shows.items():
        put(ws, c, f, F_NOTE, None, fmt, AC, border=False)
    ws.merge_cells('M3:R3')
    put(ws, 'M3', (f'=IF(AND({S("起")}>0,{S("止")}>0,{S("起")}>{S("止")}),"⚠ 起晚于止，请改日期　","")'
                   f'&IF(AND({S("日期选")}<>"",{S("日期选")}<>"预订日期",{S("日期选")}<>"出行日期"),"⚠ 按哪个日期只能选预订日期或出行日期（先按预订日期）　","")'
                   f'&IF(AND({S("状态选")}<>"",{S("状态码")}=0,{S("状态选")}<>"全部"),"⚠ 状态没认出来（先按全部）　","")'
                   f'&IF(AND({S("销售")}<>"",COUNTIFS(单_有效,1,单_销售,{esc(S("销售"))})=0),"⚠ 没有销售是「"&{S("销售")}&"」的订单","")'),
        F_RED, align=ALW, border=False)
    ws.merge_cells('M4:S4')
    put(ws, 'M4', '黄格空着＝全部；起止只填一个也行；客户填名字里的几个字就行', F_NOTE, align=AL, border=False)

    # ── 第 5～7 行：合计区（按筛出来的全部订单） ──
    lbl_fill = fill('FFD9E1F2')
    crit = f'单_有效,1,{KEY_RNG},">0"'
    kpis = [('A', 'B', '单数', f'={S("单数")}', '0'),
            ('C', 'D', '订单总金额', f'=ROUND(SUMIFS(单_订单总金额,{crit}),2)', MONEY),
            ('E', 'F', '定金金额', f'=ROUND(SUMIFS(单_定金,{crit}),2)', MONEY),
            ('G', 'H', '已收', f'=ROUND(SUMIFS(单_已收,{crit}),2)', MONEY),
            ('I', 'J', '还没收', f'=ROUND(SUMIFS(单_还没收,{crit}),2)', MONEY),
            ('K', 'L', '预计利润\n（填了预计成本的）', f'=ROUND(SUMIFS(单_预计利润,{crit},单_有预计,1),2)', MONEY),
            ('M', 'N', '实际利润\n（已出行/已结算/已取消）', f'=ROUND(SUMIFS(单_实际利润,{crit},单_出行码,">=1"),2)', MONEY),
            ('O', 'P', '订单提成\n（含预估）', f'=ROUND(SUMIFS(单_订单提成,{crit}),2)', MONEY),
            ('Q', 'R', '其中应发\n（已结算的）', f'=ROUND(SUMIFS(单_订单提成,{crit},单_应发,1),2)', MONEY)]
    for c1, c2, lbl, f, fmt in kpis:
        ws.merge_cells(f'{c1}5:{c2}5')
        ws.merge_cells(f'{c1}6:{c2}6')
        put(ws, f'{c1}5', lbl, F_KPI_L, lbl_fill, align=ACW)
        put(ws, f'{c2}5', None, F_KPI_L, lbl_fill)
        put(ws, f'{c1}6', f, F_KPI_V, FILL_TOT, fmt, AC)
        put(ws, f'{c2}6', None, F_KPI_V, FILL_TOT)
    put(ws, 'S5', '合计', F_KPI_L, lbl_fill, align=AC)
    put(ws, 'S6', '（筛出来的全部）', F_NOTE, FILL_TOT, align=ACW)
    ws.row_dimensions[5].height = 32
    ws.row_dimensions[6].height = 26
    ws.merge_cells(f'A7:{LAST}7')
    md = f'IF({S("模式")}=2,"出行日期","预订日期")'
    put(ws, 'A7', (f'=IF({S("单数")}=0,"没有符合条件的订单",'
                   f'"共 "&{S("单数")}&" 单，按"&{md}&"从早到晚排（同一天按登记顺序；没填"&{md}&"的在最后）"'
                   f'&IF({S("单数")}>{SHOW},"　⚠ 只列了前 {SHOW} 单（上面合计是全部的），请缩小日期范围","")'
                   f'&IF({S("已取消单数")}>0,"；其中已取消 "&{S("已取消单数")}&" 单（订单金额 "&TEXT({S("已取消金额")},"#,##0.00")&"）","")'
                   f'&IF({S("没填预计")}>0,"；没填预计成本 "&{S("没填预计")}&" 单（预计利润不算）","")'
                   f'&IF({S("暂算利润")}<>0,"；实际利润里暂算 "&TEXT({S("暂算利润")},"#,##0.00"),"")'
                   f'&IF({S("预估提成")}<>0,"；提成里预估 "&TEXT({S("预估提成")},"#,##0.00")&"（没结算，灰色斜体）",""))'),
        F_NOTE, align=AL, border=False)
    ws.row_dimensions[7].height = 20

    # ── 表头 ──
    heads = [('A', '序号'), ('B', '订单号'), ('C', '客户名字'), ('D', '预订日期'), ('E', '定金金额'), ('F', '出行日期'),
             ('G', '订单总金额'), ('H', '预计利润'), ('I', '实际利润\n（灰斜体＝暂算）'), ('J', '尾款收否'), ('K', '销售名字'),
             ('L', '订单提成\n（灰斜体＝预估）'), ('M', '线路 / 目的地'), ('N', '人数'), ('O', '已收'), ('P', '还没收'),
             ('Q', '实际成本\n（含未付）'), ('R', '出行状态'), ('S', '结算日期')]
    header(ws, HDR, heads[:12], C_HDR2, height=36)
    header(ws, HDR, heads[12:], C_HDR3, height=36)

    # ── 排序键（每个订单一格） ──
    ws[f'{KEY}{R0 - 1}'] = '排序键'
    ws[f'{IDX}{R0 - 1}'] = '第几条'
    ws[f'{KEY}{R0 - 1}'].font = ws[f'{IDX}{R0 - 1}'].font = F_HELP
    for i in range(N_ORD):
        n = i + 1
        u = lambda k: f'INDEX(单_{k},{n})'
        d = f'IF({S("模式")}=2,{u("出行日期")},{u("预订日期")})'
        cond = (f'OR({S("有日期条件")}=0,AND({d}>0,{d}>={S("起")},{d}<={S("止用")})),'
                f'OR({S("销售")}="",{u("销售")}={S("销售")}),'
                f'OR({S("状态码")}=0,AND({S("状态码")}<=4,{u("出行码")}={S("状态码")}-1),'
                f'AND({S("状态码")}=5,{u("出行码")}<>3,{u("还没收")}>0)),'
                f'OR({S("客户")}="",IFERROR(FIND({S("客户")},{u("客户名字")}),0)>0)')
        c = ws[f'{KEY}{R0 + i}']
        c.value = f'=IF({u("有效")}<>1,"",IF(AND({cond}),IF({d}>0,{d},99999)*10000+{n},""))'
        c.font = F_HELP

    # ── 清单 ──
    fm = {'A': '0', 'D': DATE, 'E': MONEY, 'F': DATE, 'G': MONEY, 'H': MONEY, 'I': MONEY, 'L': MONEY, 'N': INT,
          'O': MONEY, 'P': MONEY, 'Q': MONEY, 'S': DATE}
    al = {c: AR for c in 'EGHILOPQ'}
    al.update({'C': AL, 'M': AL})
    cols = {
        'D': 'IF({u}(单_预订日期,{i})>0,{u}(单_预订日期,{i}),"没填")',
        'E': '{u}(单_定金,{i})',
        'F': 'IF({u}(单_出行日期,{i})>0,{u}(单_出行日期,{i}),"没填")',
        'G': '{u}(单_订单总金额,{i})',
        'H': 'IF({u}(单_有预计,{i})=1,{u}(单_预计利润,{i}),"没填")',
        'I': '{u}(单_实际显示,{i})',
        'J': '{u}(单_尾款收否,{i})',
        'K': '{u}(单_销售,{i})',
        'L': '{u}(单_订单提成,{i})',
        'M': '{u}(单_线路,{i})',
        'N': '{u}(单_人数,{i})',
        'O': '{u}(单_已收,{i})',
        'P': '{u}(单_还没收,{i})',
        'Q': '{u}(单_实际成本,{i})',
        'R': '{u}(单_出行状态,{i})',
        'S': 'IF({u}(单_结算日期,{i})>0,{u}(单_结算日期,{i}),"")',
    }
    nn, ns = S('单数'), S('显示数')
    for j in range(NROW):
        r, k = R0 + j, j + 1
        ix = f'${IDX}{r}'
        ws[f'{IDX}{r}'] = f'=IF({k}>{ns},0,MOD(SMALL({KEY_RNG},{k}),10000))' if k <= SHOW else 0
        ws[f'{IDX}{r}'].font = F_HELP
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k})'
        ws[f'B{r}'] = f'=IF({ix}>0,INDEX(单_订单号,{ix}),IF({k}={ns}+1,IF({nn}=0,"（没有）","共 "&{nn}&" 单"),""))'
        ws[f'C{r}'] = f'=IF({ix}>0,INDEX(单_客户名字,{ix}),IF(AND({k}={ns}+1,{nn}>{SHOW}),"只列前{SHOW}单",""))'
        for col, f in cols.items():
            ws[f'{col}{r}'] = f'=IF({ix}=0,"",' + f.format(u='INDEX', i=ix) + ')'
        for col in 'ABCDEFGHIJKLMNOPQRS':
            x = ws[f'{col}{r}']
            x.font = F_TXT
            x.alignment = al.get(col, AC)
            if col in fm:
                x.number_format = fm[col]
    R1 = R0 + NROW - 1
    ix0 = f'${IDX}{R0}'
    cf_blocks(ws, R0, R1, [
        ('A', LAST, lambda c: f'AND({ix0}=0,$B{R0}<>"")', dict(font=F_GREY)),                           # 「共 N 单」
        ('A', LAST, lambda c: f'AND({ix0}>0,$R{R0}="已取消")', dict(font=F_GREY, border=CF_BD)),         # 已取消：整行灰
        ('I', 'I', lambda c: f'AND({ix0}>0,$R{R0}="已出行")', dict(font=F_GREY_I, border=CF_BD)),        # 暂算
        ('L', 'L', lambda c: f'AND({ix0}>0,$R{R0}<>"已结算")', dict(font=F_GREY_I, border=CF_BD)),       # 预估
        ('J', 'J', lambda c: f'AND({ix0}>0,$R{R0}<>"已取消",N($P{R0})>0)', dict(font=F_ORANGE_B, border=CF_BD)),  # 没收齐
        ('H', 'H', lambda c: f'AND({ix0}>0,$H{R0}="没填")', dict(font=F_GREY, border=CF_BD)),
        ('A', LAST, lambda c: f'{ix0}>0', dict(border=CF_BD)),
    ])

    hide(ws, *[CL(i) for i in range(CI(SL), CI(KEY) + 1)])
    ws.freeze_panes = f'C{R0}'
    print_setup(ws, f'{HDR}:{HDR}', landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("最后一行")})')
