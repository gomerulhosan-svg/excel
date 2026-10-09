# -*- coding: utf-8 -*-
"""【数据校验】：两张登记表逐行问题（✗ 要改、⚠ 要看）、基础资料重名、对数检查和该跟进的事，最后列出有问题的行（表名＋行号＋问题）。
   首页「提醒」引用 CHK_X / CHK_W / CHECK_ROWS。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *

CHK_X, CHK_W = 'C4', 'F4'
CHECK_ROWS = {}


def _over(sh, hdr, cap, col='A'):
    """录入表第 cap 条以后还有没有东西（区域从表头行锚定，插删行跟着动）"""
    return f'COUNTA(INDEX({sh}!${col}:${col},ROW({sh}!${col}${hdr})+{cap + 1}):INDEX({sh}!${col}:${col},1048576))'


def _over_base():
    """基础资料每块：第 cap 个以后、下一块标题以前还有没有名字（块变短时 X位<=cap，不算超）"""
    parts = []
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        a = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+{cap + 1})'
        b = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+P_{key}位)'
        parts.append(f'IF(P_{key}位<={cap},0,COUNTA({a}:{b}))')
    return '=' + '+'.join(parts)


D, Y = 'P_截止', 'P_截止年月'
OK = '单_有效,1'
OUT_KINDS = ['订单成本', '工资', '销售提成', '日常费用', '其他支出']
ARR_KIND = '{"' + '","'.join(CAT_KINDS) + '"}'


def build(wb, ctx=None):
    ws = wb[SH_CHK]
    widths(ws, {'A': 6, 'B': 50, 'C': 14, 'D': 7, 'E': 12, 'F': 14, 'G': 56})
    title(ws, '数 据 校 验', 'G', C_CHK,
          '💡 录完数看一眼这张：✗＝一定要改（不改报表会少算或算错），⚠＝看一下对不对、该跟进了，ℹ＝说明一下（不算问题）。'
          '先看每张表有几条问题，再看对数检查，最后按「表＋行号」回到那张表找到改掉。这张表全是自动的。')
    home_link(ws, 'G3')
    put(ws, 'B4', '要改的（✗）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'E4', '要看的（⚠）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    xs, wsum = [], []

    # ① 两张登记表
    r = 6
    section(ws, r, 'A', 'G', '① 两张登记表逐行检查（每行最右边「这一行的问题」汇总）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '表'), ('C', '✗ 要改'), ('D', ''), ('E', '⚠ 要看'), ('F', ''), ('G', '常见原因')], C_CHK)
    for i, (sh, nm, why) in enumerate([
            (SH_ORD, '单_校验', '没填订单号或订单号重复、没填预订日期或订单总金额、日期写成了文字、定金收款账户不在基础资料里'),
            (SH_CASH, '收_校验', '没填日期或金额、收支类别不在基础资料里、订单的钱没选订单号或订单号找不到、工资提成没选人员')]):
        rr = r + 2 + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        link(put(ws, f'B{rr}', sh, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), align=AL), sh)
        put(ws, f'C{rr}', f'=COUNTIF({nm},"✗*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'D{rr}', f'=IF(C{rr}>0,"✗","✓")', F_TXTB, align=AC)
        put(ws, f'E{rr}', f'=COUNTIF({nm},"⚠*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'F{rr}', f'=IF(E{rr}>0,"⚠","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', why, F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 30
        xs.append(f'C{rr}')
        wsum.append(f'E{rr}')
        CHECK_ROWS[sh] = rr
    r += 5

    # ② 基础资料
    section(ws, r, 'A', 'G', '② 基础资料（名字不能重复、类别要有归类）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '个数'), ('D', '状态'), ('G', '怎么改')], C_CHK)
    hcols = {'账户': ('AA', '账户_名称', N_ACC), '类别': ('AB', '类别_名称', N_CAT), '人员': ('AC', '人员_姓名', N_PER)}
    for key, (hc, rng, n) in hcols.items():
        for i in range(n):
            nm = f'INDEX({rng},{i + 1})'
            ws[f'{hc}{i + 1}'] = f'=IF({nm}="",0,IF(COUNTIF({rng},{esc(nm)})>1,1,0))'
            ws[f'{hc}{i + 1}'].font = F_HELP
    for i in range(N_CAT):                                   # 类别有名字没归类（或归类写错）
        nm, kd = f'INDEX(类别_名称,{i + 1})', f'INDEX(类别_归类,{i + 1})'
        ws[f'AD{i + 1}'] = f'=IF({nm}="",0,IF(ISNUMBER(MATCH({kd},{ARR_KIND},0)),0,1))'
        ws[f'AD{i + 1}'].font = F_HELP
    hide(ws, 'AA', 'AB', 'AC', 'AD')
    checks = [('账户名称重复', f'=SUM($AA$1:$AA${N_ACC})', '✗', '同一个账户只写一行'),
              ('收支类别名称重复', f'=SUM($AB$1:$AB${N_CAT})', '✗', '同一个类别只写一行'),
              ('人员姓名重复', f'=SUM($AC$1:$AC${N_PER})', '✗', '重名的加个字区分，比如 小李A、小李B'),
              ('收支类别没设归类（或归类不在下拉里）', f'=SUM($AD$1:$AD${N_CAT})', '✗', '在【基础资料】③ 收支类别的「归类」列从下拉选'),
              ('一个账户都没有', '=IF(COUNTIF(账户_名称,"?*")=0,1,0)', '✗', '在【基础资料】② 账户 至少填一个（比如对公账户、微信）'),
              ('订单号重复的单数', '=COUNTIF(单_校验,"✗ 订单号重复*")', '✗', '每单的订单号要不一样；第 3 行有「下一个订单号」')]
    for i, (lab, f, kind, how) in enumerate(checks):
        rr = r + 2 + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        put(ws, f'B{rr}', lab, F_TXT, align=AL)
        put(ws, f'C{rr}', f, F_AUTOB, fmt='0', align=AC)
        put(ws, f'D{rr}', f'=IF(C{rr}>0,"{kind}","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        (xs if kind == '✗' else wsum).append(f'IF(C{rr}>0,1,0)')
        CHECK_ROWS[lab] = rr
    r += len(checks) + 3

    # ③ 对数检查、该跟进的
    section(ws, r, 'A', 'G', '③ 对数检查、该跟进的（截至首页截止日期）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '金额/个数'), ('D', '状态'), ('G', '说明')], C_CHK)
    for i in range(N_ACC):                                   # 每个账户截至截止日的余额 < 0 ？
        a = f'INDEX(账户_名称,{i + 1})'
        ws[f'AE{i + 1}'] = (f'=IF({a}="",0,IF(INDEX(账户_期初余额,{i + 1})+SUMIFS(单_定金,{OK},单_收款账户,{esc(a)},'
                            f'单_定金日期,">="&P_建账日,单_定金日期,"<="&{D})+SUMIFS(收_净额,收_账户,{esc(a)})<-0.005,1,0))')
        ws[f'AE{i + 1}'].font = F_HELP
    hide(ws, 'AE')
    got = (f'SUMIFS(单_定金计入,{OK})+SUMIFS(收_已付额,收_有效,1,收_归类,"订单收款",收_日期,"<="&{D},收_订单找到,1)'
           f'+SUMIFS(收_已付额,收_有效,1,收_归类,"订单退款",收_日期,"<="&{D},收_订单找到,1)')
    soon = f'{OK},单_出行码,0,单_出行日期,">0",单_出行日期,"<="&({D}+P_尾款天数),单_还没收,">0.005"'
    late = f'{OK},单_出行码,1,单_出行日期,"<"&({D}-P_结算天数)'
    num = [
        ('内部转账没配对（转出、转入合计不为 0）', '=ROUND(SUMIFS(收_净额,收_归类,"内部转账"),2)', 'ABS(C{r})>=0.01', '✗',
         '现金存银行、微信提现要记两行（转出账户一行支出、转入账户一行收入），金额一样；用【资金流水】选类别「内部转账」查'),
        ('订单已收合计对不上（订单号重复或订单号找不到会这样）', f'=ROUND(SUMIFS(单_已收,{OK})-({got}),2)', 'ABS(C{r})>=0.01', '✗',
         '订单登记的已收合计 ≠ 定金＋收支登记里挂上订单的收款；先改订单号重复的'),
        ('订单的钱没挂上订单（订单号没选或找不到）', ('=ROUND(SUMIFS(收_支出净额,收_有效,1,收_订单找到,0,收_归类,"订单成本")'
                                         '-SUMIFS(收_支出净额,收_有效,1,收_订单找到,0,收_归类,"订单收款")'
                                         '+SUMIFS(收_支出净额,收_有效,1,收_订单找到,0,收_归类,"订单退款"),2)'),
         'ABS(C{r})>=0.01', '✗', '这些钱不进任何订单的已收、成本；到收支登记补上订单号（看最右边「这一行的问题」）'),
        ('工资、提成发给了不在人员表里的人（金额）', '=ROUND(SUMIFS(收_支出净额,收_校验,"✗ 工资、提成的往来对象*"),2)',
         'ABS(C{r})>=0.01', '✗', '往来对象要从人员下拉里选，【工资提成】才对得上人'),
        ('账户余额是负数的账户个数', f'=SUM($AE$1:$AE${N_ACC})', 'C{r}>0', '⚠', '钱不够付还付出去了？多半是漏记了收款，或者期初余额没填'),
        ('首页截止日期比最后一笔早', f'=IF(AND(ISNUMBER({SH_HOME}!{HOME_END}),P_截止<P_最后日期),P_最后日期-P_截止,0)', 'C{r}>0', '⚠',
         '首页黄格填了截止日期，以后录的就不算进来；看完了记得清空'),
        ('日期晚于截止日期的定金、收支（暂不计入）', f'=COUNTIFS({OK},单_定金,"<>0",单_定金日期,">"&{D})+COUNTIFS(收_有效,1,收_日期,">"&{D})',
         'C{r}>0', '⚠', '还没到日子的不算收到、不算付出'),
        ('快出发还没收齐的单数', f'=COUNTIFS({soon})', 'C{r}>0', '⚠', '看【未出行订单】，催尾款'),
        ('已出行还没收齐的金额', f'=ROUND(SUMIFS(单_还没收,{OK},单_出行码,">=1",单_出行码,"<=2",单_还没收,">0.005"),2)', 'C{r}>0.005', '⚠', '看【未出行订单】②'),
        ('出行后好多天还没结算的单数', f'=COUNTIFS({late})', 'C{r}>0', '⚠', '结团、成本对完账、尾款收齐后，在订单登记填结算日期（实际利润、提成才定下来）'),
        ('应退客户还没退的金额', '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单退款"),2)', 'C{r}>0.005', '⚠', '收支登记里「退客户款」付款情况是未付的；退了以后清掉「未付」、日期改成退款那天'),
        ('欠供应商的金额（订单成本未付）', '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单成本"),2)', 'C{r}>0.005', 'ℹ', '看【未出行订单】③；付了以后清掉「未付」、日期改成付款那天'),
        ('没填预计成本的单数（预估利润没算进去）', f'=COUNTIFS({OK},单_出行码,"<2",单_有预计,0)', 'C{r}>0', '⚠', '在订单登记补上预计成本'),
        ('建账日期以前的定金、收付（已含在期初余额里，不进资金）', (f'=ROUND(SUMIFS(单_定金,{OK},单_定金日期,">0",单_定金日期,"<"&P_建账日)'
                                                f'+SUMIFS(收_收入,收_有效,1,收_日期,"<"&P_建账日)+SUMIFS(收_支出,收_有效,1,收_日期,"<"&P_建账日),2)'),
         'C{r}>0', 'ℹ', '补录的老订单正常会有；这些钱订单里照算，资金余额不再算一遍'),
        (f'订单登记超过 {N_ORD} 单（多出来的不算）', '=' + _over(SH_ORD, ORD_HDR, N_ORD), 'C{r}>0', '✗', '一本最多这么多单：另起一本（没出行的单带过去）'),
        (f'收支登记超过 {N_CASH} 笔（多出来的不算）', '=' + _over(SH_CASH, CASH_HDR, N_CASH), 'C{r}>0', '✗', '同上'),
        ('基础资料某一块超了（多出来的不算）', _over_base(), 'C{r}>0', '✗', f'账户最多 {N_ACC} 个、收支类别 {N_CAT} 个、人员 {N_PER} 个；块里不要留太多空行'),
        ('基础资料没填建账日期', f'=IF(ISNUMBER({SH_BASE}!$C${PAR_ROW["建账日"]}),0,1)', 'C{r}>0', '⚠', '没填就按最早一笔算；填上更清楚（账户期初余额都是这一天的）'),
        ('基础资料没填公司名称', '=IF(P_公司名称="",1,0)', 'C{r}>0', 'ℹ', '首页、订单详情（结算单）抬头用'),
    ]
    r += 2
    for i, (lab, f, bad, kind, how) in enumerate(num):
        rr = r + i
        if lab == '快出发还没收齐的单数':
            put(ws, f'B{rr}', f'="快出发（"&P_尾款天数&" 天内）还没收齐的单数"', F_TXT, align=ALW)
        elif '好多天' in lab:
            put(ws, f'B{rr}', f'="出行超过 "&P_结算天数&" 天还没结算的单数"', F_TXT, align=ALW)
        else:
            put(ws, f'B{rr}', lab, F_TXT, align=ALW)
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        isn = any(w in lab for w in ('个数', '单数', '超', '没填', '比最后', '晚于'))
        put(ws, f'C{rr}', f, F_AUTOB, fmt='0' if isn else MONEY, align=AR)
        put(ws, f'D{rr}', f'=IF({bad.format(r=rr)},"{kind}","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 30
        if kind == '✗':
            xs.append(f'IF({bad.format(r=rr)},1,0)')
        elif kind == '⚠':
            wsum.append(f'IF({bad.format(r=rr)},1,0)')
        CHECK_ROWS[lab] = rr
    r += len(num) + 1
    put(ws, CHK_X, '=' + '+'.join(xs), F_KPI_V, fill('FFFFFFFF'), '0', AC)
    put(ws, CHK_W, '=' + '+'.join(wsum), F_KPI_V, fill('FFFFFFFF'), '0', AC)

    # ④ 有问题的行
    section(ws, r, 'A', 'G', '④ 有问题的行（回到那张表按行号找；只列 ✗ 和 ⚠）', C_CHK)
    header(ws, r + 1, [('A', '第几条'), ('B', '问题'), ('C', '表'), ('D', '行号'), ('E', '日期'), ('F', '金额'), ('G', '订单 / 摘要')], C_CHK)
    r += 2
    SHOW = 150
    lists = [(SH_ORD, N_ORD, '单_校验', '单_录入行', '单_预订日期', 'INDEX(单_订单总金额,{k})',
              'INDEX(单_订单号,{k})&"｜"&INDEX(单_客户名字,{k})', 'AJ'),
             (SH_CASH, N_CASH, '收_校验', '收_录入行', '收_日期', 'INDEX(收_收入,{k})-INDEX(收_支出,{k})',
              'INDEX(收_类别,{k})&"｜"&INDEX(收_显示摘要,{k})', 'AK')]
    for sh, n, chk, row_nm, date_nm, amt, desc, cc in lists:
        counter(ws, cc, 1, n, lambda i, chk=chk: f'OR(LEFT(INDEX({chk},{i + 1}),1)="✗",LEFT(INDEX({chk},{i + 1}),1)="⚠")')
        hide(ws, cc)
        total = f'${cc}${n}'
        put(ws, f'B{r}', f'="{sh}：共 "&{total}&" 行有问题"&IF({total}>{SHOW},"（只列前 {SHOW} 行）","")', F_TXTB, FILL_SUB, align=AL)
        ws.merge_cells(f'B{r}:G{r}')
        r += 1
        for k in range(1, SHOW + 1):
            rr = r + k - 1
            ix = f'$A{rr}'
            ws[f'A{rr}'] = f'=IF({kth(k, cc, 1, n)}=0,"",{kth(k, cc, 1, n)})'
            ws[f'A{rr}'].font = F_HELP
            g = lambda e: f'=IF({ix}="","",{e.format(k=ix)})'
            put(ws, f'B{rr}', g(f'INDEX({chk},{{k}})'), F_TXT, align=AL)
            put(ws, f'C{rr}', f'=IF({ix}="","","{sh}")', F_TXT, align=AC)
            put(ws, f'D{rr}', g(f'INDEX({row_nm},{{k}})'), F_TXTB, fmt='0', align=AC)
            put(ws, f'E{rr}', g(f'IF(INDEX({date_nm},{{k}})=0,"",INDEX({date_nm},{{k}}))'), F_TXT, fmt=DATE, align=AC)
            put(ws, f'F{rr}', g(amt), F_TXT, fmt=MONEY, align=AR)
            put(ws, f'G{rr}', g(desc), F_NOTE, align=AL)
        r += SHOW + 1
    for col in ('D', 'F'):
        ws.conditional_formatting.add(f'{col}1:{col}{r}', FormulaRule(formula=[f'{col}1="✗"'], font=F_RED))
        ws.conditional_formatting.add(f'{col}1:{col}{r}', FormulaRule(formula=[f'{col}1="⚠"'],
                                                                     font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.conditional_formatting.add(f'B1:B{r}', FormulaRule(formula=['LEFT(B1,1)="✗"'], font=F_RED))
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:4', landscape=False)
    ws.print_area = f'A1:G{r}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
