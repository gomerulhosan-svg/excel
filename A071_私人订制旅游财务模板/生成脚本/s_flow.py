# -*- coding: utf-8 -*-
"""查看表（绿）【资金流水】出纳的现金 / 银行日记账：订单的定金（订单登记）＋ 收支登记里已付的钱，按日期排，带余额。
   黄格：B3 起（空＝建账日；早于建账日按建账日）、D3 止（空＝截止日；晚于截止日按截止日）、F3 账户（空＝全部）、H3 类别（归类；空＝全部）。
   ① 各账户：期初（起日前一天）、本期收入、本期支出、期末；合计行（账户名不在基础资料里的差额单列一行）。① 不受类别影响。
      期初＝账户_期初余额＋[建账日, 起) 的定金和 收_净额；收入＝[起, 止] 的定金（正的）＋收_收入（收_在期=1）；支出＝收_支出（收_在期=1）。
   ② 流水：定金（单_有效=1、单_定金<>0、单_定金日期 在 [起, 止]）和收支登记 收_在期=1 的行，按 日期×SORT_M＋id 排（同一天收支在前、定金在后）。
      余额只在类别空时算：选了账户＝这个账户的余额，没选＝全部账户合计，从①的期初往下滚；合计行核对期末＝①的期末。
   排序键在隐藏表 _序 的 A 列（第 2 行起）：id 1..N_CASH＝收支第 n 条，N_CASH+1..N_CASH+N_ORD＝订单第 n 条的定金。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'J'
SHOW = 2000
KINDS = ['定金'] + CAT_KINDS                           # 类别下拉：定金＋收支类别的归类
R_SEC1, R_H1, R_A0 = 6, 7, 8
NA = N_ACC + 2                                         # ① 最多 N_ACC 个账户＋「账户名不对」＋合计
R_SEC2 = R_A0 + NA + 1
R_NOTE2, R_H2 = R_SEC2 + 1, R_SEC2 + 2
R_D0 = R_H2 + 1
R_D1 = R_D0 + SHOW + 1                                 # 明细最后一行（第 SHOW+2 行：合计、共几笔）

# ── 隐藏列 ──
SL, SV = 'AA', 'AB'
SC = ['账户', '类别', '账户在', '类别对', '起', '止', '账户数', '笔数', '显示数', '定金笔数', '期初', '本期收入', '本期支出', '期末算',
      '账户位', '①期末', '滚到', '核对', '末行', '其他要显示', '止算']
# 止算＝MAX(止, 起－1)：起晚于止（日期范围是空的）时，期末＝期初，①、② 的核对照样对得上（第 5 行另有红字提示改日期）
SR = {k: 3 + i for i, k in enumerate(SC)}
HR0 = 5                                                # ① 每个账户的数（第 i 个在 HR0+i-1 行；之后「账户名不对」、合计）
HA = dict(名称='AD', 期初='AE', 收入='AF', 支出='AG', 期末='AH')
A_OTH, A_TOT = N_ACC + 1, N_ACC + 2
IDX1 = 'AI'                                            # ① 显示行取第几个
HID, HKD, HN, HGL = 'AJ', 'AK', 'AL', 'AM'             # ② id、类型（1 收支 2 定金 3 合计 4 共几笔）、第几条、归类
NKEY = N_CASH + N_ORD
TAB_ACC = f'{H_TAB}!${COMPACT["账户"][0]}$2:${COMPACT["账户"][0]}${N_ACC + 1}'
TAB_ACC_N = f'{H_TAB}!${COMPACT["账户"][0]}$1'

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
GREEN_H = 'FF70AD47'
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC00000')


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def R2(x):
    return f'ROUND({x},2)'


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


TIP = ('💡 出纳的现金、银行日记账：订单的定金（订单登记填的）和收支登记里已经付了的钱，按日期排，余额逐笔往下滚，可以拿来跟银行流水、'
       '微信账单、现金盘点逐笔对。黄格：起、止（空着＝建账日～首页截止日）、账户（空＝全部账户合起来）、类别（空＝全部；定金、订单收款……）。'
       '① 每个账户的期初（起日前一天）、本期收入、支出、期末余额（不受类别影响）；② 逐笔流水，选了类别时只列这一类、不算余额。'
       '「未付」的还没付出去，不在这里；建账日以前的钱已经含在账户期初余额里。现金存银行、微信提现（内部转账）在各账户里是真进真出，'
       '选全部账户时一出一进、余额不变。要改哪一笔，按最后一列行号去【订单登记】或【收支登记】改。')


def build(wb, ctx=None):
    ws = wb[SH_FLOW]
    widths(ws, {'A': 11.5, 'B': 10, 'C': 13, 'D': 11, 'E': 14, 'F': 26, 'G': 13.5, 'H': 13.5, 'I': 14.5, 'J': 10})
    title(ws, '资金流水（现金 / 银行日记账）', LAST, C_VIEW, TIP)

    # ── 选择格 ──
    selector(ws, 'A3', '起', 'B3', None, fmt=DATE, prompt='空着＝建账日')
    selector(ws, 'C3', '止', 'D3', None, fmt=DATE, prompt='空着＝首页截止日；晚于截止日按截止日')
    selector(ws, 'E3', '账户', 'F3', None, '=账户列表', prompt='从下拉选账户；空着＝全部账户合起来')
    selector(ws, 'G3', '类别', 'H3', None, '"' + ','.join(KINDS) + '"', prompt='只看这一类的钱（按收支类别的归类）；空着＝全部')
    dv_date(ws, 'B3')
    dv_date(ws, 'D3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 26

    ACC, CAT, S0, S1 = S('账户'), S('类别'), S('起'), S('止')
    S1X = S('止算')
    accN = any_crit('账户_名称', ACC)
    accD = any_crit('单_收款账户', ACC)
    accS = any_crit('收_账户', ACC)
    catS = any_crit('收_归类', CAT)
    dep_ok = f'OR({CAT}="",{CAT}="定金")'
    kinds_txt = ',' + ','.join(KINDS) + ','
    dep_before = lambda a: f'SUMIFS(单_定金,单_有效,1,{a},单_定金,"<>0",单_定金日期,">="&P_建账日,单_定金日期,"<"&{S0},单_定金日期,"<="&P_截止)'
    dep_upto = lambda a: f'SUMIFS(单_定金,单_有效,1,{a},单_定金,"<>0",单_定金日期,">="&P_建账日,单_定金日期,"<="&{S1X},单_定金日期,"<="&P_截止)'
    dep_in = lambda a: f'SUMIFS(单_定金,单_有效,1,{a},单_定金,">0",{dr("单_定金日期", S0, S1)})'
    dep_out = lambda a: f'-SUMIFS(单_定金,单_有效,1,{a},单_定金,"<0",{dr("单_定金日期", S0, S1)})'
    cash_before = lambda a: f'SUMIFS(收_净额,收_有效,1,{a},收_日期,"<"&{S0})'
    cash_upto = lambda a: f'SUMIFS(收_净额,收_有效,1,{a},收_日期,"<="&{S1X})'
    cash_in = lambda a, c='': f'SUMIFS(收_收入,收_在期,1,收_有效,1,{a}{c},{dr("收_日期", S0, S1)})'
    cash_out = lambda a, c='': f'SUMIFS(收_支出,收_在期,1,收_有效,1,{a}{c},{dr("收_日期", S0, S1)})'
    KEYS = f"{H_SORT}!$A$2:$A${NKEY + 1}"
    lst = lambda c: f'${c}${R_D0}:${c}${R_D0 + SHOW - 1}'
    eA = esc(ACC)
    sc = {
        '账户': '=TRIM(F3&"")',
        '类别': '=TRIM(H3&"")',
        '账户在': f'=IF({ACC}="",1,IF(COUNTIF(账户_名称,{eA})>0,1,0))',
        '类别对': f'=IF({CAT}="",1,IF(ISNUMBER(FIND(","&{CAT}&",","{kinds_txt}")),1,0))',
        '起': '=IF(ISNUMBER(B3),MAX(INT(B3),P_建账日),P_建账日)',
        '止': '=IF(ISNUMBER(D3),MIN(INT(D3),P_截止),P_截止)',
        '账户数': f'=MIN(N({TAB_ACC_N}),{N_ACC})',
        '笔数': f'=COUNT({KEYS})',
        '显示数': f'=MIN({S("笔数")},{SHOW})',
        '定金笔数': f'=COUNT({H_SORT}!$A${N_CASH + 2}:$A${NKEY + 1})',
        '期初': '=' + R2(f'SUMIFS(账户_期初余额,{accN})+{dep_before(accD)}+{cash_before(accS)}'),
        '本期收入': '=' + R2(f'IF({dep_ok},{dep_in(accD)},0)+{cash_in(accS, "," + catS)}'),
        '本期支出': '=' + R2(f'IF({dep_ok},{dep_out(accD)},0)+{cash_out(accS, "," + catS)}'),
        '期末算': f'=ROUND({S("期初")}+{S("本期收入")}-{S("本期支出")},2)',
        '账户位': f'=IF({ACC}="",0,IFERROR(MATCH({eA},{TAB_ACC},0),0))',
        '①期末': (f'=IF({ACC}="",${HA["期末"]}${HR0 + A_TOT - 1},IF({S("账户位")}>0,'
                 f'INDEX(${HA["期末"]}${HR0}:${HA["期末"]}${HR0 + N_ACC - 1},{S("账户位")}),'
                 + R2(f'SUMIFS(账户_期初余额,账户_名称,{eA})+{dep_upto("单_收款账户," + eA)}+{cash_upto("收_账户," + eA)}') + '))'),
        '滚到': f'=IF({S("显示数")}=0,{S("期初")},N(INDEX({lst("I")},{S("显示数")})))',
        '核对': (f'=IF({CAT}<>"",-1,IF(AND(ROUND({S("滚到")}-{S("①期末")},2)=0,ROUND({S("期末算")}-{S("①期末")},2)=0),1,0))'),
        '末行': f'={R_D0 + 1}+{S("显示数")}',
        '其他要显示': '=IF(OR(' + ','.join(f'ROUND({HA[k]}{HR0 + A_OTH - 1},2)<>0' for k in ('期初', '收入', '支出', '期末')) + '),1,0)',
        '止算': f'=MAX({S1},{S0}-1)',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    # ── 第 4、5 行：实际用的 ──
    put(ws, 'A4', '实际用的', F_NOTE, align=AC, border=False)
    put(ws, 'B4', f'={S0}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'D4', f'={S1}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'F4', f'=IF({ACC}="","全部账户",{ACC})', F_AUTOB, FILL_AUTO, align=AC)
    put(ws, 'H4', f'=IF({CAT}="","全部类别",{CAT})', F_AUTOB, FILL_AUTO, align=AC)
    put(ws, 'C4', '', F_NOTE, border=False)
    ws.merge_cells(f'A5:{LAST}5')
    put(ws, 'A5', (f'=IF({S0}>{S1},IF(AND(ISNUMBER(D3),INT(N(D3))<P_建账日),"⚠ 止早于建账日（建账以前的钱已含在期初余额里），请改日期　",'
                   f'"⚠ 起晚于止，这段时间是空的（期末＝期初），请改日期　"),"")'
                   f'&IF(AND(TRIM(B3&"")<>"",NOT(ISNUMBER(B3))),"⚠ 起不是日期（先按建账日算）　","")'
                   f'&IF(AND(TRIM(D3&"")<>"",NOT(ISNUMBER(D3))),"⚠ 止不是日期（先按截止日算）　","")'
                   f'&IF(AND(ISNUMBER(B3),INT(N(B3))<P_建账日),"起早于建账日，按建账日 "&TEXT(P_建账日,"yyyy/mm/dd")&" 算（建账以前的钱已含在期初余额里）　","")'
                   f'&IF(AND(ISNUMBER(D3),INT(N(D3))>P_截止),"止晚于截止日，按截止日 "&TEXT(P_截止,"yyyy/mm/dd")&" 算（截止日在首页改）　","")'
                   f'&IF({S("账户在")}=0,"⚠ 「"&{ACC}&"」不在【基础资料】的账户里　","")'
                   f'&IF({S("类别对")}=0,"⚠ 类别「"&{CAT}&"」不对，请从下拉选","")'), F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add('A5', FormulaRule(formula=['ISNUMBER(FIND("⚠",$A$5))'], font=F_WARN))
    ws.row_dimensions[4].height = 20

    # ═════ ① 各账户（隐藏表 AD:AH，第 HR0 行起） ═════
    for c, t in zip(HA.values(), HA.keys()):
        ws[f'{c}{HR0 - 1}'] = t
        ws[f'{c}{HR0 - 1}'].font = F_HELP
    for i in range(1, N_ACC + 1):
        r = HR0 + i - 1
        a = f'${HA["名称"]}{r}'
        e = esc(a)
        f = {
            '名称': f'=IF({i}>{S("账户数")},"",INDEX({TAB_ACC},{i})&"")',
            '期初': f'=IF({a}="",0,{R2(f"SUMIFS(账户_期初余额,账户_名称,{e})+" + dep_before("单_收款账户," + e) + "+" + cash_before("收_账户," + e))})',
            '收入': f'=IF({a}="",0,{R2(dep_in("单_收款账户," + e) + "+" + cash_in("收_账户," + e))})',
            '支出': f'=IF({a}="",0,{R2(dep_out("单_收款账户," + e) + "+" + cash_out("收_账户," + e))})',
            '期末': f'=IF({a}="",0,{R2(f"SUMIFS(账户_期初余额,账户_名称,{e})+" + dep_upto("单_收款账户," + e) + "+" + cash_upto("收_账户," + e))})',
        }
        for k, v in f.items():
            ws[f'{HA[k]}{r}'] = v
            ws[f'{HA[k]}{r}'].font = F_HELP
    ro, rt = HR0 + A_OTH - 1, HR0 + A_TOT - 1
    ws[f'{HA["名称"]}{ro}'] = '账户名不对'
    ws[f'{HA["名称"]}{rt}'] = '合计'
    allA = '账户_名称,"<>@@全部@@"'
    tot = {
        '期初': R2(f'SUMIFS(账户_期初余额,{allA})+{dep_before("单_有效,1")}+{cash_before("收_有效,1")}'),
        '收入': R2(f'{dep_in("单_有效,1")}+{cash_in("收_有效,1")}'),
        '支出': R2(f'{dep_out("单_有效,1")}+{cash_out("收_有效,1")}'),
        '期末': R2(f'SUMIFS(账户_期初余额,{allA})+{dep_upto("单_有效,1")}+{cash_upto("收_有效,1")}'),
    }
    for k, v in tot.items():
        c = HA[k]
        ws[f'{c}{rt}'] = '=' + v
        ws[f'{c}{ro}'] = f'=ROUND({c}{rt}-SUM({c}{HR0}:{c}{HR0 + N_ACC - 1}),2)'
        ws[f'{c}{rt}'].font = ws[f'{c}{ro}'].font = F_HELP

    section(ws, R_SEC1, 'A', LAST, '① 各账户（期初＝起日前一天的余额；收入含定金；不受类别影响；选的账户黄底）', C_VIEW)
    header(ws, R_H1, [('A', '账户'), ('C', '期初余额\n（起日前一天）'), ('D', '本期收入'), ('E', '本期支出'), ('F', '期末余额'),
                      ('G', '说明')], GREEN_H, height=36)
    ws.merge_cells(f'A{R_H1}:B{R_H1}')
    ws.merge_cells(f'G{R_H1}:{LAST}{R_H1}')
    ws[f'B{R_H1}'].border = BD
    for c in 'HIJ':
        ws[f'{c}{R_H1}'].border = BD
    hr = lambda col: f'${col}${HR0}:${col}${HR0 + A_TOT - 1}'
    na, oth = S('账户数'), S('其他要显示')
    for k in range(1, NA + 1):
        r = R_A0 + k - 1
        z = f'${IDX1}{r}'
        ws[f'{IDX1}{r}'] = (f'=IF({k}<={na},{k},IF({k}={na}+1,IF({oth}=1,{A_OTH},{A_TOT}),'
                            f'IF(AND({k}={na}+2,{oth}=1),{A_TOT},0)))')
        ws[f'{IDX1}{r}'].font = F_HELP
        ws.merge_cells(f'A{r}:B{r}')
        for c, key in (('A', '名称'), ('C', '期初'), ('D', '收入'), ('E', '支出'), ('F', '期末')):
            ws[f'{c}{r}'] = f'=IF({z}=0,"",INDEX({hr(HA[key])},{z}))'
        ws.merge_cells(f'G{r}:{LAST}{r}')
        ws[f'G{r}'] = (f'=IF({z}=0,"",IF({z}<={N_ACC},IF(AND({ACC}<>"",A{r}={ACC}),"◀ 选的账户　","")'
                       f'&IF(F{r}<0,"⚠ 余额是负数：漏记了收入，还是记错了账户？",""),'
                       f'IF({z}={A_OTH},"⚠ 有账户名不在【基础资料】的账户里（看订单登记、收支登记最右边的校验列）",'
                       f'IF(ROUND(C{r}+D{r}-E{r}-F{r},2)=0,"✓ 期初＋收入－支出＝期末","⚠ 期初＋收入－支出≠期末"))))')
        for c in 'ABCDEFGHIJ':
            x = ws[f'{c}{r}']
            x.font = F_NOTE if c in 'GHIJ' else F_TXT
            x.alignment = AL if c in 'ABGHIJ' else AR
            if c in 'CDEF':
                x.number_format = MONEY
    a0, a1 = R_A0, R_A0 + NA - 1
    rng1 = f'A{a0}:{LAST}{a1}'
    z0 = f'${IDX1}{a0}'
    cf = ws.conditional_formatting
    cf.add(f'G{a0}:{LAST}{a1}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",$G{a0}))'], font=Font(bold=True, color='FFC00000')))
    cf.add(rng1, FormulaRule(formula=[f'{z0}={A_TOT}'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng1, FormulaRule(formula=[f'{z0}={A_OTH}'], font=Font(italic=True, color='FFC00000'), border=CF_BD, stopIfTrue=True))
    cf.add(rng1, FormulaRule(formula=[f'AND({z0}>0,{ACC}<>"",$A{a0}={ACC})'], fill=cf_fill('FFFFF2CC'), border=CF_BD,
                             stopIfTrue=True))
    cf.add(rng1, FormulaRule(formula=[f'{z0}>0'], border=CF_BD))

    # ═════ ② 流水 ═════
    section(ws, R_SEC2, 'A', LAST, '② 流水（定金＋收支登记里已付的钱，按日期排；余额从①的期初逐笔往下滚）', C_VIEW)
    N, NN = S('笔数'), S('显示数')
    ws.merge_cells(f'A{R_NOTE2}:{LAST}{R_NOTE2}')
    put(ws, f'A{R_NOTE2}', (f'=IF({N}=0,"这段时间没有流水"&IF({CAT}="","（余额不变：期初＝期末）",""),'
                            f'"共 "&{N}&" 笔（其中定金 "&{S("定金笔数")}&" 笔），按日期排（同一天先列收支登记、后列定金）")'
                            f'&IF({N}>{SHOW},"　⚠ 只显示前 {SHOW} 笔（余额也只滚到第 {SHOW} 笔），请缩短日期范围","")'
                            f'&IF({CAT}="","；余额＝"&IF({ACC}="","全部账户合计","「"&{ACC}&"」")&"的余额，从期初 "'
                            f'&TEXT({S("期初")},"#,##0.00")&" 往下滚","；选了类别：只列这一类，不算余额")'),
        F_NOTE, align=AL, border=False)
    cf.add(f'A{R_NOTE2}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",$A${R_NOTE2}))'], font=F_WARN))
    header(ws, R_H2, [('A', '日期'), ('B', '账户'), ('C', '类别'), ('D', '订单号'), ('E', '客户 / 往来对象'), ('F', '摘要'),
                      ('G', '收入'), ('H', '支出'), ('I', '余额'), ('J', '登记行号')], GREEN_H, height=30)
    for c, t in ((HID, 'id'), (HKD, '类型'), (HN, '第几条'), (HGL, '归类')):
        ws[f'{c}{R_H2}'] = t
        ws[f'{c}{R_H2}'].font = F_HELP
    open_ = S('期初')
    for k in range(1, SHOW + 3):
        r = R_D0 + k - 1
        y, q, n = f'${HID}{r}', f'${HKD}{r}', f'${HN}{r}'
        ws[f'{HID}{r}'] = f'=IF({k}>{NN},0,MOD(SMALL({KEYS},{k}),{SORT_M}))' if k <= SHOW else 0
        ws[f'{HKD}{r}'] = f'=IF({y}>0,IF({y}<={N_CASH},1,2),IF({k}={NN}+1,3,IF({k}={NN}+2,4,0)))'
        ws[f'{HN}{r}'] = f'=IF({q}=1,{y},IF({q}=2,{y}-{N_CASH},0))'
        ws[f'{HGL}{r}'] = f'=IF({q}=1,INDEX(收_归类,{n}),IF({q}=2,"定金",""))'
        for c in (HID, HKD, HN, HGL):
            ws[f'{c}{r}'].font = F_HELP
        s_ = lambda f_: f'INDEX(收_{f_},{n})'
        d_ = lambda f_: f'INDEX(单_{f_},{n})'
        prev = open_ if k == 1 else f'N(I{r - 1})'
        chk = (f'IF({CAT}<>"","选了类别：不算余额",IF({N}>{SHOW},"⚠ 只列了前 {SHOW} 笔（合计行的收入、支出、余额是全期的）",'
               f'IF({S("核对")}=1,"✓ 期末余额＝①的期末","✗ 跟①的期末 "&TEXT({S("①期末")},"#,##0.00")&" 对不上")))')
        cells = {
            'A': [(f'{q}=1', s_('日期')), (f'{q}=2', d_('定金日期')), (f'{q}=3', '"合计"'), (f'{q}=4', f'"共 "&{N}&" 笔"')],
            'B': [(f'{q}=1', s_('账户')), (f'{q}=2', d_('收款账户'))],
            'C': [(f'{q}=1', f'IF({s_("类别")}="","（没选类别）",{s_("类别")})'), (f'{q}=2', '"定金"')],
            'D': [(f'{q}=1', s_('订单号')), (f'{q}=2', d_('订单号'))],
            'E': [(f'{q}=1', f'IF({s_("往来对象")}<>"",{s_("往来对象")},IF({s_("订单行")}>0,INDEX(单_客户名字,{s_("订单行")}),""))'),
                  (f'{q}=2', d_('客户名字'))],
            'F': [(f'{q}=1', s_('摘要')), (f'{q}=2', f'"收定金"&IF({d_("线路")}<>"","（"&{d_("线路")}&"）","")'), (f'{q}=3', chk)],
            'G': [(f'{q}=1', f'IF({s_("收入")}=0,"",{s_("收入")})'), (f'{q}=2', f'IF({d_("定金")}>0,{d_("定金")},"")'),
                  (f'{q}=3', S('本期收入'))],
            'H': [(f'{q}=1', f'IF({s_("支出")}=0,"",{s_("支出")})'), (f'{q}=2', f'IF({d_("定金")}<0,-{d_("定金")},"")'),
                  (f'{q}=3', S('本期支出'))],
            'I': [(f'OR({q}=1,{q}=2)', f'IF({CAT}="",ROUND({open_}+SUM(G${R_D0}:G{r})-SUM(H${R_D0}:H{r}),2),"")'),   # 不成链：WPS 改一格不用整列重算
                  (f'{q}=3', f'IF({CAT}="",IF({N}>{SHOW},{S("①期末")},{prev}),"")')],    # 超出：合计行给真的期末（收入、支出也是全期的）
            'J': [(f'{q}=1', f'"收支 "&{s_("录入行")}'), (f'{q}=2', f'"订单 "&{d_("录入行")}')],
        }
        for c, pairs in cells.items():
            x = ws[f'{c}{r}']
            x.value = '=' + _chain(pairs)
            x.font = F_TXT
            x.alignment = AL if c in 'BCDEF' else (AR if c in 'GHI' else AC)
            if c == 'A':
                x.number_format = DATE
            elif c in 'GH':
                x.number_format = MONEY_B
            elif c == 'I':
                x.number_format = MONEY
    d0, d1 = R_D0, R_D1
    rng2 = f'A{d0}:{LAST}{d1}'
    q0 = f'${HKD}{d0}'
    cf.add(f'F{d0}:F{d1}', FormulaRule(formula=[f'AND({q0}=3,OR(LEFT($F{d0},1)="✗",LEFT($F{d0},1)="⚠"))'],
                                       font=Font(bold=True, color='FFC00000')))
    cf.add(rng2, FormulaRule(formula=[f'{q0}=3'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng2, FormulaRule(formula=[f'{q0}=4'], font=Font(color='FF808080'), stopIfTrue=True))
    cf.add(rng2, FormulaRule(formula=[f'AND({q0}=1,${HGL}{d0}="内部转账")'], fill=cf_fill('FFDDEBF7'), border=CF_BD,
                             stopIfTrue=True))
    cf.add(rng2, FormulaRule(formula=[f'{q0}=2'], fill=cf_fill('FFF4F9EE'), border=CF_BD, stopIfTrue=True))
    cf.add(rng2, FormulaRule(formula=[f'{q0}=1'], border=CF_BD))

    # ── 隐藏表 _序：排序键（日期×SORT_M＋id；不在期间/账户/类别里的、金额为 0 的放空） ──
    hs = wb[H_SORT]
    hs['A1'] = '资金流水排序键'
    hs['B1'] = 'id'
    hs['C1'] = '第几条'
    q = f"'{SH_FLOW}'!"
    s0, s1, acc, cat = (q + S(k) for k in ('起', '止', '账户', '类别'))
    for i in range(NKEY):
        r = i + 2
        b, cc = f'$B{r}', f'$C{r}'
        hs[f'B{r}'] = i + 1
        if i < N_CASH:
            hs[f'C{r}'] = i + 1
            x = lambda f_: f'INDEX(收_{f_},{cc})'
            hs[f'A{r}'] = (f'=IF({x("在期")}<>1,"",IF(AND({x("日期")}>={s0},{x("日期")}<={s1},OR({acc}="",{x("账户")}={acc}),'
                           f'OR({cat}="",{x("归类")}={cat})),{x("日期")}*{SORT_M}+{b},""))')
        else:
            hs[f'C{r}'] = i + 1 - N_CASH
            x = lambda f_: f'INDEX(单_{f_},{cc})'
            hs[f'A{r}'] = (f'=IF({x("有效")}<>1,"",IF(AND({x("定金")}<>0,{x("定金日期")}>=MAX(P_建账日,{s0}),'
                           f'{x("定金日期")}<=MIN(P_截止,{s1}),OR({acc}="",{x("收款账户")}={acc}),OR({cat}="",{cat}="定金")),'
                           f'{x("定金日期")}*{SORT_M}+{b},""))')

    hide(ws, *[CL(i) for i in range(CI('K'), CI(HGL) + 1)])
    ws.freeze_panes = 'A5'
    print_setup(ws, f'{R_H2}:{R_H2}', landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    qq = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{qq}!$A$1:INDEX({qq}!${LAST}$1:${LAST}${d1},{qq}!{S("末行")})')
    return ws
