# -*- coding: utf-8 -*-
"""查看表（绿）【月度汇总】选一个年份：1～12 月 × 预订、回款、支出、其他、结算、经营利润（估）、出行；下面一个快照框（截至截止日，不分月）。
   黄格 B3 年份（空＝截止日那年；填 2026、"2026"、或者那年里的一个日期都行；认不出来按截止日那年并红字提醒）。
   A 列项目名（用户原话做名字，括号里写口径），B～M 列 1～12 月，N 列全年（＝1～12 月相加）。
   口径（agent_common「口径速查」）：
   - 预订按预订月（含后来取消的）；当月取消按取消月（单_出行码=3）；
   - 钱按收付日期，只算建账日～截止日、已付的：定金（单_定金日期 在 [P_建账日, P_截止]）＋ 收_净额 按 收_归类 分；
   - 结算（单_应发=1）按结算月：单数、订单金额、成本、利润、应发提成；取消单留存利润按取消月；
   - 经营利润（估）＝结算利润＋取消单留存＋其他收入－月工资－销售提成（应发）－日常费用－其他支出；
   - 出行按出行月，不含取消。
   快照框：全部预计订单未出行预估总利润（单_出行码=0、单_有预计=1 的预计利润）、未出行单数/金额/已收/还没收/预计提成；
   没填预计成本的单数和金额；已出行没结算；欠供应商、应退客户、其他未付（收_未付额 按归类）。
   隐藏列：AA 说明、AB 值（第 3 行起）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from layout import *
from common import *

LAST = 'N'
HDR, R0 = 5, 6
MCOL = [CL(2 + i) for i in range(12)]                   # B～M：1～12 月
SL, SV = 'AA', 'AB'
SC = ['年份输入', '年份', '年份认出', '截止年月',
      '预估总利润', '未出行单数', '未出行金额', '未出行已收', '未出行还没收', '未出行提成', '没预计单数', '没预计金额',
      '待结单数', '待结预计利润', '待结暂算利润', '欠供应商', '应退客户', '未付合计', '其他未付']
SR = {k: 3 + i for i, k in enumerate(SC)}

LBL = fill('FFD9E1F2')
F_BIG = Font(name=YH, sz=20, bold=True, color='FFC00000')
F_BIGL = Font(name=YH, sz=12, bold=True, color='FF1F3864')
F_GREY_I = Font(name=YH, sz=10, italic=True, color='FF9E9E9E')
F_GREY_IB = Font(name=YH, sz=10, bold=True, italic=True, color='FF9E9E9E')
F_LINK = Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single')
CNT = '0" 单";-0" 单";"0 单"'


def S(k):
    return f'${SV}${SR[k]}'


def R2(x):
    return f'ROUND({x},2)'


def K(kind, y, neg=False):
    """收支登记里这一归类、这个月的净额（在期：已付、建账日～截止日）；neg＝支出类取反成正数"""
    return R2(f'{"-" if neg else ""}SUMIFS(收_净额,收_有效,1,收_归类,"{kind}",收_年月,{y})')


DEP = lambda y: R2(f'SUMIFS(单_定金,单_有效,1,单_定金年月,{y},单_定金日期,">="&P_建账日,单_定金日期,"<="&P_截止)')
ST = lambda y: f'单_有效,1,单_应发,1,单_结算年月,{y}'           # 结算（按结算月）
CX = lambda y: f'单_有效,1,单_出行码,3,单_取消年月,{y}'          # 取消（按取消月）
TR = lambda y: f'单_有效,1,单_出行年月,{y},单_出行码,"<>3"'      # 出行（按出行月，不含取消）
BK = lambda y: f'单_有效,1,单_预订年月,{y}'                      # 预订（按预订月）


def ref(*keys, sign=None):
    """本列里别的行相加减：ref('a','b','c', sign='++-')"""
    sign = sign or '+' * len(keys)
    return lambda y, c, R: R2(''.join(f'{"" if i == 0 and s == "+" else s}{c}{R[k]}' for i, (k, s) in enumerate(zip(keys, sign))))


# ('S', 分区标题) 或 (键, 项目名, 样式, 公式)；样式：n＝单数，''＝金额，key＝用户原话（加粗、FILL_SUB），tot＝经营利润
# 公式 f(y, c, R)：y＝这个月 YYYYMM 的表达式，c＝本列字母，R＝键→行号
ITEMS = [
    ('S', '一、预订（按预订日期）'),
    ('预订单数', '预订单数', 'n', lambda y, c, R: f'COUNTIFS({BK(y)})'),
    ('月预订金额', '月预订金额（订单总金额，含后来取消的）', 'key', lambda y, c, R: R2(f'SUMIFS(单_订单总金额,{BK(y)})')),
    ('预订预计利润', '其中预计利润（填了预计成本的；已取消的按留存）', '', lambda y, c, R: R2(f'SUMIFS(单_预计利润,{BK(y)},单_有预计,1)')),
    ('取消单数', '当月取消单数（按取消日期）', 'n', lambda y, c, R: f'COUNTIFS({CX(y)})'),
    ('取消金额', '当月取消金额（订单总金额）', '', lambda y, c, R: R2(f'SUMIFS(单_订单总金额,{CX(y)})')),
    ('S', '二、回款（按收到的日期；只算建账日～截止日）'),
    ('定金', '定金（订单登记里填的）', '', lambda y, c, R: DEP(y)),
    ('收尾款', '收尾款/分期', '', lambda y, c, R: K('订单收款', y)),
    ('退款', '退客户款（减）', '', lambda y, c, R: K('订单退款', y, True)),
    ('月回款金额', '月回款金额（＝定金＋收尾款/分期－退客户款）', 'key', ref('定金', '收尾款', '退款', sign='++-')),
    ('S', '三、支出（按付款日期；未付的不算）'),
    ('订单成本', '订单成本（已付：地接、机票、酒店、税费等）', '', lambda y, c, R: K('订单成本', y, True)),
    ('月工资', '月工资', 'key', lambda y, c, R: K('工资', y, True)),
    ('已发提成', '销售提成（已发）', 'key', lambda y, c, R: K('销售提成', y, True)),
    ('日常费用', '日常费用（房租、推广、办公等）', '', lambda y, c, R: K('日常费用', y, True)),
    ('其他支出', '其他支出', '', lambda y, c, R: K('其他支出', y, True)),
    ('月支出金额', '月支出金额（＝上面 5 项合计）', 'key', ref('订单成本', '月工资', '已发提成', '日常费用', '其他支出')),
    ('S', '四、其他收入、净进账'),
    ('其他收入', '其他收入（季度返佣、利息等）', '', lambda y, c, R: K('其他收入', y)),
    ('净进账', '本月净进账（＝月回款＋其他收入－月支出）', '', ref('月回款金额', '其他收入', '月支出金额', sign='++-')),
    ('资金往来', '资金往来净额（股东投入取出、借还款；不算收支）', '', lambda y, c, R: K('资金往来', y)),
    ('S', '五、结算（按结算日期：结团对完账的单）'),
    ('结算单数', '结算单数', 'n', lambda y, c, R: f'COUNTIFS({ST(y)})'),
    ('结算金额', '结算订单金额', '', lambda y, c, R: R2(f'SUMIFS(单_订单总金额,{ST(y)})')),
    ('结算成本', '结算订单成本（含未付）', '', lambda y, c, R: R2(f'SUMIFS(单_实际成本,{ST(y)})')),
    ('结算利润', '结算利润（实际）', '', lambda y, c, R: R2(f'SUMIFS(单_实际利润,{ST(y)})')),
    ('取消留存', '取消单留存利润（按取消日期）', '', lambda y, c, R: R2(f'SUMIFS(单_实际利润,{CX(y)})')),
    ('应发提成', '销售提成（应发：当月结算的单）', 'key', lambda y, c, R: R2(f'SUMIFS(单_订单提成,{ST(y)})')),
    ('S', '六、经营利润（估）＝结算利润＋取消单留存＋其他收入－月工资－销售提成（应发）－日常费用－其他支出'),
    ('经营利润', '经营利润（估）', 'tot', ref('结算利润', '取消留存', '其他收入', '月工资', '应发提成', '日常费用', '其他支出',
                                       sign='+++----')),
    ('S', '七、出行（按出行日期，不含取消）'),
    ('出行单数', '出行单数', 'n', lambda y, c, R: f'COUNTIFS({TR(y)})'),
    ('出行金额', '出行订单金额', '', lambda y, c, R: R2(f'SUMIFS(单_订单总金额,{TR(y)})')),
]

TIP = ('💡 全自动不用填。黄格 B3 填年份（比如 2026；空着＝截止日期那年），出 1～12 月和全年。用户要的五个数——月预订金额、月回款金额、'
       '月支出金额、月工资、销售提成——是加粗浅蓝的行。预订按预订日期（含后来取消的）；回款、支出按钱真正收到/付出的日期，只算建账日～截止日、'
       '已付的（「未付」的成本不算支出）；销售提成分「已发」（收支登记里发出去的）和「应发」（当月结算的单该发的）；结算利润、经营利润（估）按结算月。'
       '灰色表头＝截止日期以后的月份（只有出行是已经排好的）。最下面是截至截止日期的未出行订单、没结算的、欠人家的钱（不分月）。金额显示到元。')


def build(wb, ctx=None):
    ws = wb[SH_MONTH]
    widths(ws, dict({'A': 30, LAST: 12}, **{c: 11 for c in MCOL}))
    title(ws, '月度汇总（每月预订、回款、支出、工资、提成、利润）', LAST, C_VIEW, TIP)

    # ── 第 3 行：年份黄格 ＋ 灰字口径 ──
    selector(ws, 'A3', '年份', 'B3', None, fmt='0')
    dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2099', allow_blank=True,
                        showErrorMessage=False, showInputMessage=True, promptTitle='提示',
                        prompt='填年份，比如 2026；空着＝截止日期那年')
    ws.add_data_validation(dv)
    dv.add('B3')
    ws.merge_cells('C3:E3')
    put(ws, 'C3', (f'="＝ "&{S("年份")}&" 年"&IF(TRIM(B3&"")="","（空着＝截止日那年）","")'
                   f'&IF({S("年份认出")}=0,"（⚠ 没认出来，先按截止日那年）","")'), F_NOTE, align=AL, border=False)
    ws.merge_cells('F3:M3')
    put(ws, 'F3', '口径：钱按收付日期，只算建账日～截止日、已付的（未付的不算）；预订按预订日期；利润、应发提成按结算月；出行按出行日期。',
        F_NOTE, align=ALW, border=False)
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 30
    ws.conditional_formatting.add('C3', FormulaRule(formula=[f'{S("年份认出")}=0'], font=F_RED))

    # ── 隐藏标量 ──
    yv = S('年份输入')
    U0 = '单_有效,1,单_出行码,0'
    U1 = '单_有效,1,单_出行码,1'
    sc = {
        '年份输入': '=IFERROR(--TRIM(B3&""),0)',
        '年份': f'=IF(AND({yv}>=2000,{yv}<=2099),INT({yv}),IF(AND({yv}>=36526,{yv}<=73050),YEAR({yv}),YEAR(P_截止)))',
        '年份认出': f'=IF(TRIM(B3&"")="",1,IF(OR(AND({yv}>=2000,{yv}<=2099),AND({yv}>=36526,{yv}<=73050)),1,0))',
        '截止年月': '=P_截止年月',
        '预估总利润': f'=ROUND(SUMIFS(单_预计利润,{U0},单_有预计,1),2)',
        '未出行单数': f'=COUNTIFS({U0})',
        '未出行金额': f'=ROUND(SUMIFS(单_订单总金额,{U0}),2)',
        '未出行已收': f'=ROUND(SUMIFS(单_已收,{U0}),2)',
        '未出行还没收': f'=ROUND(SUMIFS(单_还没收,{U0}),2)',
        '未出行提成': f'=ROUND(SUMIFS(单_订单提成,{U0}),2)',
        '没预计单数': f'=COUNTIFS({U0},单_有预计,0)',
        '没预计金额': f'=ROUND(SUMIFS(单_订单总金额,{U0},单_有预计,0),2)',
        '待结单数': f'=COUNTIFS({U1})',
        '待结预计利润': f'=ROUND(SUMIFS(单_预计利润,{U1}),2)',
        '待结暂算利润': f'=ROUND(SUMIFS(单_实际利润,{U1}),2)',
        '欠供应商': '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单成本"),2)',
        '应退客户': '=ROUND(SUMIFS(收_未付额,收_有效,1,收_归类,"订单退款"),2)',
        '未付合计': '=ROUND(SUMIFS(收_未付额,收_有效,1),2)',
        '其他未付': f'=ROUND({S("未付合计")}-{S("欠供应商")}-{S("应退客户")},2)',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    # ── 第 4 行：截止日、建账日 ──
    ws.merge_cells(f'A4:{LAST}4')
    put(ws, 'A4', ('="截止日期 "&TEXT(P_截止,"yyyy-mm-dd")&"（【首页】改），建账日 "&TEXT(P_建账日,"yyyy-mm-dd")'
                   '&"：回款、支出只算这段时间里已付的，建账以前收付的钱在账户期初余额里。"'
                   f'&IF({S("年份")}<YEAR(P_建账日),"⚠ 选的年份在建账以前：回款、支出都是 0（预订、结算、出行照算）。","")'
                   f'&IF({S("年份")}>YEAR(P_截止),"⚠ 选的年份在截止日期以后：只有已经排好的出行。","")'),
        F_NOTE, align=AL, border=False)
    ws.row_dimensions[4].height = 20
    ws.conditional_formatting.add('A4', FormulaRule(formula=[f'OR({S("年份")}<YEAR(P_建账日),{S("年份")}>YEAR(P_截止))'],
                                                    font=Font(color='FFC65911')))

    # ── 第 5 行表头 ──
    header(ws, HDR, [('A', f'="项目（"&{S("年份")}&" 年）"')] + [(c, f'{m}月') for m, c in enumerate(MCOL, 1)] + [(LAST, '全年')],
           C_VIEW, height=24)
    grey_h = PatternFill('solid', fgColor='FFA6A6A6', bgColor='FFA6A6A6')
    for m, c in enumerate(MCOL, 1):
        ws.conditional_formatting.add(f'{c}{HDR}', FormulaRule(formula=[f'{S("年份")}*100+{m}>{S("截止年月")}'], fill=grey_h))

    # ── 表身 ──
    R = {}
    r = R0
    rows = []
    for it in ITEMS:
        rows.append((r, it))
        if it[0] != 'S':
            R[it[0]] = r
        r += 1
    last_item = r - 1
    for r, it in rows:
        if it[0] == 'S':
            section(ws, r, 'A', LAST, it[1], C_VIEW)
            ws.row_dimensions[r].height = 20
            continue
        key, lbl, sty, fn = it
        fmt = INT if sty == 'n' else MONEY0
        bold = sty in ('key', 'tot')
        fl = FILL_SUB if sty == 'key' else (FILL_TOT if sty == 'tot' else None)
        put(ws, f'A{r}', lbl, F_TXTB if bold else F_TXT, fl, align=ALW)
        for m, c in enumerate(MCOL, 1):
            y = f'({S("年份")}*100+{m})'
            put(ws, f'{c}{r}', '=' + fn(y, c, R), F_AUTOB if bold else F_AUTO, fl, fmt, AR)
        put(ws, f'{LAST}{r}', f'=SUM({MCOL[0]}{r}:{MCOL[-1]}{r})', F_AUTOB, fl or FILL_AUTO, fmt, AR)
        ws.row_dimensions[r].height = 30 if len(lbl) > 14 else 20

    # ── 快照框（截至截止日，不分月） ──
    rs = last_item + 2
    section(ws, rs, 'A', LAST, f'="八、截至 "&TEXT(P_截止,"yyyy-mm-dd")&" 的情况（不分月，跟上面选的年份无关）"', C_VIEW)
    ws.row_dimensions[rs].height = 20
    a, b = rs + 1, rs + 2
    put(ws, f'A{a}', '全部预计订单未出行预估总利润', F_BIGL, LBL, align=ACW)
    put(ws, f'A{b}', f'={S("预估总利润")}', F_BIG, FILL_TOT, MONEY0, AC)
    kp = [('B', 'C', '未出行单数', S('未出行单数'), CNT, F_AUTOB),
          ('D', 'E', '订单金额', S('未出行金额'), MONEY0, F_AUTOB),
          ('F', 'G', '已收', S('未出行已收'), MONEY0, F_AUTOB),
          ('H', 'I', '还没收', S('未出行还没收'), MONEY0, F_AUTOB),
          ('J', 'K', '预计提成（预估）', S('未出行提成'), MONEY0, F_GREY_IB)]
    for c1, c2, lbl, f, fmt, fnt in kp:
        ws.merge_cells(f'{c1}{a}:{c2}{a}')
        ws.merge_cells(f'{c1}{b}:{c2}{b}')
        put(ws, f'{c1}{a}', lbl, F_KPI_L, LBL, align=ACW)
        put(ws, f'{c2}{a}', None, F_KPI_L, LBL)
        put(ws, f'{c1}{b}', f'={f}', fnt, FILL_AUTO, fmt, AC)
        put(ws, f'{c2}{b}', None, fnt, FILL_AUTO)
    ws.merge_cells(f'L{a}:{LAST}{b}')
    put(ws, f'L{a}', '预估总利润＝还没出行（没取消）的订单「订单总金额－预计成本」合计；没填预计成本的不算', F_NOTE, align=ALW)
    for rr in (a, b):
        for c in ('L', 'M', LAST):
            ws[f'{c}{rr}'].border = BD
    ws.row_dimensions[a].height = 30
    ws.row_dimensions[b].height = 34
    c3 = rs + 3
    ws.merge_cells(f'A{c3}:{LAST}{c3}')
    put(ws, f'A{c3}', (f'=IF({S("没预计单数")}=0,"未出行的订单都填了预计成本 ✓",'
                       f'"另有 "&{S("没预计单数")}&" 单（订单金额 "&TEXT({S("没预计金额")},"#,##0.00")'
                       f'&"）没填预计成本，没算进上面的预估总利润——去【订单登记】补上预计成本")'), F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'A{c3}', FormulaRule(formula=[f'{S("没预计单数")}>0'], font=Font(color='FFC65911')))
    ws.row_dimensions[c3].height = 18
    a, b = rs + 4, rs + 5
    ws.merge_cells(f'A{a}:A{b}')
    put(ws, f'A{a}', '没结算的、欠人家的', F_KPI_L, LBL, align=ACW)
    put(ws, f'A{b}', None, F_KPI_L, LBL)
    kp = [('B', 'C', '已出行没结算', S('待结单数'), CNT, F_AUTOB),
          ('D', 'E', '其预计利润', S('待结预计利润'), MONEY0, F_AUTOB),
          ('F', 'G', '其暂算实际利润', S('待结暂算利润'), MONEY0, F_GREY_IB),
          ('H', 'I', '欠供应商（成本未付）', S('欠供应商'), MONEY0, F_AUTOB),
          ('J', 'K', '应退客户（退款未付）', S('应退客户'), MONEY0, F_AUTOB),
          ('L', 'M', '其他未付', S('其他未付'), MONEY0, F_AUTOB)]
    for c1, c2, lbl, f, fmt, fnt in kp:
        ws.merge_cells(f'{c1}{a}:{c2}{a}')
        ws.merge_cells(f'{c1}{b}:{c2}{b}')
        put(ws, f'{c1}{a}', lbl, F_KPI_L, LBL, align=ACW)
        put(ws, f'{c2}{a}', None, F_KPI_L, LBL)
        put(ws, f'{c1}{b}', f'={f}', fnt, FILL_AUTO, fmt, AC)
        put(ws, f'{c2}{b}', None, fnt, FILL_AUTO)
    ws.merge_cells(f'{LAST}{a}:{LAST}{b}')
    link(put(ws, f'{LAST}{a}', '看明细 →', F_LINK, fill('FFFFFFFF'), align=ACW), SH_TODO)
    ws[f'{LAST}{b}'].border = BD
    ws.row_dimensions[a].height = 30
    ws.row_dimensions[b].height = 24
    ws.conditional_formatting.add(f'B{b}:C{b}', FormulaRule(formula=[f'{S("待结单数")}>0'], font=Font(bold=True, color='FFC00000')))
    last = b

    hide(ws, SL, SV)
    ws.freeze_panes = f'B{R0}'
    print_setup(ws, f'{HDR}:{HDR}', landscape=True)
    ws.print_area = f'A1:{LAST}{last}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
