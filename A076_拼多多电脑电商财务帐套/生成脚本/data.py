# -*- coding: utf-8 -*-
"""读用户上传的原件（参考/上传原件_拼多多帐套资料.xlsx）→ 导入帐套的数据。原样保留管家婆、拼多多导出的写法（文字金额、文字日期都不改）。"""
import datetime as dt
import openpyxl


def _clean(row, n):
    row = list(row[:n]) + [None] * max(0, n - len(row))
    out = []
    for v in row:
        if isinstance(v, str):
            if v == '':
                v = None
            elif len(v) > 32000:
                v = v[:32000]
        out.append(v)
    return out


def load(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    pur = [_clean(r, 18) for r in wb['采购'].iter_rows(min_row=6, values_only=True) if isinstance(r[0], int)]
    sal = [_clean(r, 13) for r in wb['销售明细'].iter_rows(min_row=6, values_only=True) if isinstance(r[0], int)]
    inv = [r for r in wb['库存'].iter_rows(min_row=7, values_only=True) if isinstance(r[0], int)]
    p1 = [_clean(r, 7) for r in wb['拼多多店铺帐务明细查询'].iter_rows(min_row=6, values_only=True)]
    while p1 and all(v is None for v in p1[-1]):
        p1.pop()
    p2 = [_clean(r, 12) for r in wb['拼多多平台营销活动结算对帐单'].iter_rows(min_row=2, values_only=True) if r[0]]
    p3 = [_clean(r, 12) for r in wb['技术服务费'].iter_rows(min_row=2, values_only=True) if r[0]]

    # 往来单位：供应商（采购单位）＋客户（拼多多买家、销售单位）
    sup, cus = [], ['拼多多买家']
    for r in pur:
        if r[3] and r[3] not in sup:
            sup.append(r[3])
    for r in sal:
        if r[1] and r[1] not in cus:
            cus.append(r[1])
    units = []
    for s in sup:
        units.append((s, '供应商+客户' if s in cus else '供应商', None))
    for c in cus:
        if c in sup:
            continue
        online = c == '拼多多买家' or '拼多多' in c
        units.append((c, '客户', '线上' if online else '线下'))
    units = [(n, t, ch if ch else ('线下' if '客户' in t else None)) for n, t, ch in units]

    # 商品：库存表 57 个（期初）＋采购、销售里出现的新商品（期初 0）
    goods, seen = [], set()
    for r in inv:
        nm = r[7]
        if nm in seen:
            continue
        seen.add(nm)
        goods.append((nm, r[10] or r[2], r[9] or None, r[13] or 0, round(float(r[15] or 0), 4)))
    for nm, br, md in [(r[8], r[11], r[10]) for r in pur] + [(r[7], None, None) for r in sal]:
        if nm and nm not in seen:
            seen.add(nm)
            goods.append((nm, br, md, None, None))

    # 资金台帐：账务明细里的提现 → 内部转账（拼多多店铺余额 → 招商银行）
    cash = []
    for r in p1:
        if r[4] == '提现' and r[1] is not None:
            t = r[1]
            cash.append((dt.datetime(t.year, t.month, t.day), '拼多多店铺余额', '内部转账', None, None, '拼多多店铺提现到招商银行',
                         None, abs(float(r[3])), '招商银行', f'账务明细 {t:%Y-%m-%d %H:%M} 提现申请'))
    return dict(pur=pur, sal=sal, p1=p1, p2=p2, p3=p3, units=units, goods=goods, cash=cash, year=2026, month=9,
                inv_total=sum(float(r[17] or 0) for r in inv))
