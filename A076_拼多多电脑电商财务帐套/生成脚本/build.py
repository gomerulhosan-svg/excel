# -*- coding: utf-8 -*-
"""A076 拼多多电脑电商财务帐套 · 生成脚本

跑法：python3 -I build.py        （-I 也能跑：本目录自己加进 sys.path）
输出：../A076-1_录入本.xlsx、../A076-2_报表本.xlsx

做法：make.build() 一本版（录入表＋报表，报表公式写 '录入表'!区域）→ LibreOffice 整本重算（关多线程），数报错、对关键合计
      → 录入本＝删报表表；报表本＝删录入表、公式里 '录入表'! 改成 '[1]录入表'!、加外部链接（目标＝相对文件名 A076-1_录入本.xlsx）
      → 两本都把一本版算好的值写回当缓存；报表本的外部链接缓存按录入本写好（单独打开报表本也先显示上次的数）。成品不经 LibreOffice 存盘。
"""
import os
import re
import shutil
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import openpyxl
from share import share
from inject_cache import inject
from lo_recalc import recalc, unhide_cols
from refresh_link_cache import refresh
import data
import make

ROOT = os.path.dirname(HERE)
SRC = os.path.join(ROOT, '参考', '上传原件_拼多多帐套资料.xlsx')
OUT1 = os.path.join(ROOT, make.WB1)
OUT2 = os.path.join(ROOT, make.WB2)
CALC = os.path.join(HERE, '_calc')

BANNED = re.compile(r'(?<![A-Z0-9_.])(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|LET|LAMBDA|MAXIFS|MINIFS|IFS|SWITCH|'
                    r'TEXTJOIN|CONCAT|INDIRECT|OFFSET|TODAY|NOW|RAND|AGGREGATE|TEXTSPLIT)\(')


def check_formulas(wb):
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    out = ''.join(v.split('"')[0::2]).upper()
                    if BANNED.findall(out) or out.count('(') != out.count(')') or len(v) > 8000:
                        bad.append(f'{ws.title}!{c.coordinate}: {v[:120]}')
    assert not bad, '\n'.join(bad[:20])


def to_report(wb):
    """一本版 → 报表本（就地）：删录入表和名称，公式里录入表加 [1]，加外部链接"""
    for n in make.ENTRY:
        del wb[n]
    for n in list(wb.defined_names):
        del wb.defined_names[n]
    n_f = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    for e in make.ENTRY:
                        v = v.replace(f"'{e}'!", f"'[1]{e}'!")
                    if v != c.value:
                        n_f += 1
                    c.value = v
                    for e in make.ENTRY:
                        assert f"'{e}'!" not in v and f'{e}!' not in v.replace(f'[1]{e}', ''), (ws.title, c.coordinate, v[:100])
    from openpyxl.workbook.external_link.external import (ExternalLink, ExternalBook, ExternalSheetNames,
                                                           ExternalSheetDataSet, ExternalSheetData)
    from openpyxl.packaging.relationship import Relationship
    eb = ExternalBook(sheetNames=ExternalSheetNames(sheetName=list(make.ENTRY)),
                      sheetDataSet=ExternalSheetDataSet(sheetData=[ExternalSheetData(sheetId=i) for i in range(len(make.ENTRY))]),
                      id='rId1')
    el = ExternalLink(externalBook=eb)
    el.file_link = Relationship(type='externalLinkPath', Target=make.WB1, TargetMode='External', Id='rId1')
    wb._external_links.append(el)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == make.HOME
    return n_f


def to_entry(wb):
    for n in make.REPORT:
        del wb[n]
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == make.BS


def key_checks(calc):
    """必要核查：采购合计、9 月店铺净额 vs #总计、批次（9/15 五笔补贴汇入、9/24 返还）"""
    wb = openpyxl.load_workbook(calc, data_only=True)
    out = []
    pur = sum(v for (v,) in wb[make.PU].iter_rows(min_row=4, max_row=2003, min_col=22, max_col=22, values_only=True) if isinstance(v, (int, float)))
    out.append(('采购合计', round(pur, 2), 2813917.82))
    inv = wb[make.INV]
    out.append(('库存表采购金额合计', round(inv['F5'].value, 2), 2813917.82))
    rec = wb[make.REC]
    rows = {rec[f'A{r}'].value: r for r in range(1, 40)}
    sep = 'K'                                               # 9 月在 C..N 的第 9 列
    out.append(('9月店铺净额', rec[f'{sep}{rows["本月净额（上面合计）"]}'].value, 61390.86))
    out.append(('9月#总计', rec[f'{sep}{rows["账务明细「#总计」"]}'].value, 61390.86))
    out.append(('9月差异', rec[f'{sep}{rows["差异（应为 0）"]}'].value, 0))
    out.append(('9月提现差异', rec[f'{sep}{rows["提现差异（台帐－账务明细）"]}'].value, 0))
    sub = ret = 0
    lines = []
    for r in range(1, rec.max_row + 1):
        c = rec[f'C{r}'].value
        if c in ('补贴汇入', '技术服务费返还') and isinstance(rec[f'E{r}'].value, (int, float)):
            lines.append(f"    {rec[f'B{r}'].value} {c} {rec[f'D{r}'].value} {rec[f'E{r}'].value} 对账单{rec[f'F{r}'].value or 0}+{rec[f'G{r}'].value or 0} → {rec[f'J{r}'].value}")
            if c == '补贴汇入':
                sub += rec[f'E{r}'].value
            elif abs(rec[f'E{r}'].value - 4586.81) < 0.001:
                ret = rec[f'G{r}'].value
    out.append(('批次A：补贴汇入到账合计', round(sub, 2), 3970))
    out.append(('批次A：9/24 返还批次对账单合计', round(ret or 0, 2), 4586.81))
    ok = True
    for k, v, want in out:
        good = isinstance(v, (int, float)) and abs(v - want) < 0.005
        ok &= good
        print(f'   {"✓" if good else "✗"} {k}：{v}（应为 {want}）')
    print('\n'.join(lines))
    tf = wb[make.TF]
    print('   技术服务费：', {tf[f'E{r}'].value: tf[f'H{r}'].value for r in range(5, 13)},
          {tf[f'B{r}'].value: tf[f'C{r}'].value for r in range(5, 12)})
    home = wb[make.HOME]
    for r in range(16, 25):
        print('   首页', home[f'A{r}'].value, home[f'B{r}'].value, home[f'C{r}'].value)
    return ok


def main():
    d = data.load(SRC)
    os.makedirs(CALC, exist_ok=True)
    one = os.path.join(CALC, 'A076_一本版.xlsx')
    wb, _ = make.build(d)
    check_formulas(wb)
    wb.save(one)
    share(one)
    calc = os.path.join(CALC, 'A076_一本版_算数.xlsx')
    shutil.copy(one, calc)
    unhide_cols(calc)
    st = recalc(calc, 3000)
    if 'error' in st:
        print('✗ 重算没跑成：', st['error']); return 2
    print(f"① 一本版 LibreOffice 重算：公式 {st['formulas']:,} 个，报错 {st['errors']} 处，{st['wall_seconds']} 秒")
    for k, v in st['error_cells'].items():
        print(f'    {k}: {v[:15]}')
    ok = key_checks(calc) and st['errors'] == 0

    wb, _ = make.build(d)
    to_entry(wb)
    wb.save(OUT1)
    share(OUT1)
    n1, _ = inject(OUT1, calc)
    wb, _ = make.build(d)
    nf = to_report(wb)
    wb.save(OUT2)
    share(OUT2)
    n2, _ = inject(OUT2, calc)
    nl = refresh(OUT2, {make.WB1: OUT1}, log=lambda *a: None)
    print(f'② 录入本 {os.path.getsize(OUT1) / 1e6:.1f} MB（缓存 {n1:,} 格）；报表本 {os.path.getsize(OUT2) / 1e6:.1f} MB'
          f'（外部引用公式 {nf:,} 个，缓存 {n2:,} 格，链接缓存 {nl:,} 格）')
    return 0 if ok else 1


if __name__ == '__main__':
    sys.exit(main())
