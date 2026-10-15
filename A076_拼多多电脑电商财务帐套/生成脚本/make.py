# -*- coding: utf-8 -*-
"""A076 拼多多电脑电商财务帐套：先在一本里把「录入表 + 报表」全建好（报表公式写 '录入表名'!区域），
   build.py 再拆成 A076-1_录入本 / A076-2_报表本（报表本里把 '录入表名'! 改成 '[1]录入表名'!，加外部链接）。
   公式只用 Excel 2010 / WPS 都有的函数；OFFSET 只用在下拉名称里。"""
import openpyxl
from openpyxl.styles import NamedStyle, Font, PatternFill, Border, Side, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI

ENTRY = ['基础资料', '资金台帐', '采购入库', '销售出库', '粘贴_账务明细', '粘贴_营销活动', '粘贴_服务费返还']
REPORT = ['首页', '拼多多对账', '技术服务费核对', '库存', '往来', '费用统计', '利润表', '使用说明']
BS, CA, PU, SA, P1, P2, P3 = ENTRY
HOME, REC, TF, INV, AR, EXP, PL, HELP = REPORT
WB1 = 'A076-1_录入本.xlsx'
WB2 = 'A076-2_报表本.xlsx'

# 行范围
CR = (4, 3003)      # 资金台帐
PR = (4, 2003)      # 采购、销售
XR = (4, 3003)      # 三张拼多多粘贴区
ACC = (5, 10)       # 资金账户 6 个；第 4 个（第 8 行）固定是拼多多店铺余额
CAT = (14, 43)      # 收支类别
MAP = (47, 61)      # 拼多多账务类型对照
FEE = (5, 54)       # 费用项目
UNIT = (5, 404)     # 往来单位
GOOD = (5, 304)     # 商品
PDD = f"'{BS}'!$A$8"
SLOTS = 'OPQRST'    # 资金台帐 6 个账户的累计余额（隐藏列）

YH = '微软雅黑'
MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
QTY = '#,##0.##;[Red]-#,##0.##;"-"'
DATE = 'yyyy-mm-dd'
DTIME = 'yyyy-mm-dd hh:mm'


def X(sh, rng):
    return f"'{sh}'!{rng}"


def RG(sh, col, r=None):
    """录入表某列的整段数据区：'表'!$G$4:$G$3003"""
    r0, r1 = r or {CA: CR, PU: PR, SA: PR}.get(sh, XR)
    return X(sh, f'${col}${r0}:${col}${r1}')


# ─────────────────────────── 样式 ───────────────────────────
def _styles(wb):
    thin = Side(style='thin', color='FFBFBFBF')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    al = Alignment(vertical='center')

    def ns(name, font=None, fill=None, fmt='General', align=None, border=True):
        s = NamedStyle(name=name)
        s.font = font or Font(name=YH, sz=10)
        if fill:
            s.fill = PatternFill('solid', fgColor=fill)
        s.number_format = fmt
        s.alignment = align or al
        if border:
            s.border = bd
        wb.add_named_style(s)
    au = Font(name=YH, sz=10, color='FF404040')
    ns('t', Font(name=YH, sz=14, bold=True, color='FFFFFFFF'), 'FF1F4E79', border=False)
    ns('tip', Font(name=YH, sz=9, color='FF7F6000'), 'FFFFF2CC', align=Alignment(vertical='center', wrap_text=True), border=False)
    ns('sec', Font(name=YH, sz=11, bold=True, color='FF1F4E79'), border=False)
    ns('h', Font(name=YH, sz=10, bold=True, color='FFFFFFFF'), 'FF2F75B5', align=Alignment(horizontal='center', vertical='center', wrap_text=True))
    ns('hh', Font(name=YH, sz=8, color='FF808080'), 'FFE7E6E6', align=Alignment(horizontal='center', vertical='center', wrap_text=True))
    for k, f in (('', 'General'), ('m', MONEY), ('d', DATE), ('n', QTY), ('t', '@')):
        ns('i' + k, fill='FFFFF7E0', fmt=f)
    for k, f in (('', 'General'), ('m', MONEY), ('d', DATE), ('dt', DTIME), ('n', QTY), ('i', '0')):
        ns('a' + k, au, 'FFF2F2F2', fmt=f)
    ns('tot', Font(name=YH, sz=10, bold=True), 'FFFCE4D6')
    ns('totm', Font(name=YH, sz=10, bold=True), 'FFFCE4D6', fmt=MONEY)
    ns('lb', Font(name=YH, sz=10))
    ns('lbw', Font(name=YH, sz=10), align=Alignment(vertical='center', wrap_text=True))
    ns('kpi', Font(name=YH, sz=12, bold=True, color='FFC00000'), 'FFF2F2F2', fmt=MONEY)
    ns('warn', Font(name=YH, sz=10, bold=True, color='FFC00000'), 'FFF2F2F2', align=Alignment(vertical='center', wrap_text=True))
    ns('lnk', Font(name=YH, sz=10, color='FF0563C1', underline='single'), border=False)
    ns('txt', Font(name=YH, sz=10), align=Alignment(vertical='top', wrap_text=True), border=False)


def put(ws, a, v=None, s=None):
    c = ws[a]
    if v is not None:
        c.value = v
    if s:
        c.style = s
    return c


def col(ws, letter, r0, r1, tpl, s):
    ci = CI(letter)
    for r in range(r0, r1 + 1):
        c = ws.cell(row=r, column=ci)
        if tpl is not None:
            c.value = tpl.replace('{r}', str(r)).replace('{p}', str(r - 1))
        c.style = s


def head(ws, row, cols, title, tip, last=None):
    """cols: [(列, 表头, 宽度)]"""
    last = last or cols[-1][0]
    ws.merge_cells(f'A1:{last}1')
    put(ws, 'A1', title, 't')
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f'A2:{last}2')
    put(ws, 'A2', tip, 'tip')
    ws.row_dimensions[2].height = 48
    for L, t, w in cols:
        put(ws, f'{L}{row}', t, 'h')
        ws.column_dimensions[L].width = w
    ws.row_dimensions[row].height = 30


def dv_list(ws, src, rng, allow_blank=True):
    dv = DataValidation(type='list', formula1=src, allow_blank=allow_blank, showErrorMessage=False)
    dv.add(rng)
    ws.add_data_validation(dv)


def red_if(ws, rng, formula):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=Font(name=YH, bold=True, color='FFC00000'),
                                                   fill=PatternFill('solid', fgColor='FFFFC7CE')))


def green_if(ws, rng, formula):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[formula], font=Font(name=YH, color='FF006100'),
                                                   fill=PatternFill('solid', fgColor='FFE2EFDA')))


# ─────────────────────────── 录入本：基础资料 ───────────────────────────
CATS = [('销售收款', '往来', '线下客户回款（选往来单位）'), ('销售退款', '往来', '退给客户（填支出）'),
        ('采购付款', '往来', '付给供应商（选往来单位）'), ('采购退款', '往来', '供应商退回（填收入）'),
        ('借款/还款', '往来', '股东、个人借还款（往来单位里先加上）'),
        ('费用支出', '支出', '日常费用（选费用项目）'), ('工资社保', '支出', ''), ('税费', '支出', ''),
        ('其他收入', '收入', '利息、返利等'), ('内部转账', '内部转账', '账户之间互转、拼多多提现（填对方账户）')]
MAPS = [('0010002', '订单收入', '店铺回款'), ('0010005', '优惠券结算', '店铺回款'), ('0020002', '交易退款', '店铺回款'),
        ('0020005', '退款扣优惠券', '店铺回款'), ('0030002', '基础技术服务费', '技术服务费'),
        ('0030003', '百亿补贴技术服务费', '技术服务费'), ('0030001', '技术服务费返还', '技术服务费'),
        ('0090001', '补贴汇入', '平台补贴'), ('0050002', '消费者体验提升计划', '体验计划'),
        ('0040002', '售后补偿', '售后补偿'), ('0130005', '申诉补回', '售后补偿'), ('0080001', '提现', '提现')]
FEES = ['房租', '水电物业', '工资', '社保公积金', '快递物流', '包装耗材', '多多推广/广告', '办公费', '通讯网络',
        '交通差旅', '业务招待', '维修费', '银行手续费', '税费', '其他']


def s_base(wb, d):
    ws = wb[BS]
    head(ws, 4, [('A', '账户名称', 16), ('B', '类型', 12), ('C', '期初余额', 14), ('D', '当前余额', 15), ('E', '备注', 30),
                 ('F', '', 2), ('G', '费用项目', 16), ('H', '', 2),
                 ('I', '往来单位名称', 30), ('J', '类型', 12), ('K', '渠道', 8), ('L', '期初应付', 12), ('M', '期初应收', 12),
                 ('N', '备注', 12), ('O', '', 2),
                 ('P', '商品名称', 44), ('Q', '品牌', 8), ('R', '型号', 14), ('S', '期初数量', 10), ('T', '期初单价', 12),
                 ('U', '期初金额', 14)],
         'A076 拼多多电脑电商 · 基础资料（录入本）',
         '淡黄格手填、灰格自动。资金账户第 4 行固定是「拼多多店铺余额」（账务明细粘进来自动算余额，提现到招行在资金台帐记内部转账）。'
         '商品期初＝上传的管家婆库存表（你说不太准），盘点后直接改「期初数量／期初单价」。往来单位：供应商、客户都在这一张表（类型区分；拼多多卖的填「线上」）。'
         '拼多多账务类型对照按「业务描述」前 7 位代码归类，名称别改（报表按名称取数），新类型往下面空行加。')
    for r, (nm, tp, note) in enumerate([('招商银行', '银行', ''), ('微信', '第三方支付', ''), ('支付宝', '第三方支付', ''),
                                        ('拼多多店铺余额', '平台余额', '货款在平台；期初填 9/1 店铺余额'), (None, None, None),
                                        (None, None, None)], ACC[0]):
        put(ws, f'A{r}', nm, 'i'); put(ws, f'B{r}', tp, 'i'); put(ws, f'C{r}', 0 if nm else None, 'im'); put(ws, f'E{r}', note, 'i')
        k = SLOTS[r - ACC[0]]
        extra = f"+SUMIFS({RG(P1, 'K')},{RG(P1, 'L')},\"<>提现\")" if r == 8 else ''
        put(ws, f'D{r}', f"=IF(A{r}=\"\",\"\",'{CA}'!{k}{CR[1]}{extra})", 'am')
    put(ws, 'A11', '合计', 'tot'); put(ws, 'C11', f'=SUM(C{ACC[0]}:C{ACC[1]})', 'totm'); put(ws, 'D11', f'=SUM(D{ACC[0]}:D{ACC[1]})', 'totm')
    put(ws, 'A12', '收支类别（属性：收入／支出／内部转账／往来）', 'sec')
    for L, t in zip('ABC', ('收支类别', '属性', '说明')):
        put(ws, f'{L}13', t, 'h')
    for i, r in enumerate(range(CAT[0], CAT[1] + 1)):
        v = CATS[i] if i < len(CATS) else (None, None, None)
        put(ws, f'A{r}', v[0], 'i'); put(ws, f'B{r}', v[1], 'i'); put(ws, f'C{r}', v[2], 'i')
    dv_list(ws, '"收入,支出,内部转账,往来"', f'B{CAT[0]}:B{CAT[1]}')
    put(ws, 'A45', '拼多多账务类型对照（账务明细「业务描述」前 7 位代码）', 'sec')
    for L, t in zip('ABC', ('业务描述代码', '归到', '利润表归类')):
        put(ws, f'{L}46', t, 'h')
    for i, r in enumerate(range(MAP[0], MAP[1] + 1)):
        v = MAPS[i] if i < len(MAPS) else (None, None, None)
        put(ws, f'A{r}', v[0], 'it'); put(ws, f'B{r}', v[1], 'i'); put(ws, f'C{r}', v[2], 'i')
    dv_list(ws, '"店铺回款,技术服务费,平台补贴,体验计划,售后补偿,其他收入,其他费用,提现"', f'C{MAP[0]}:C{MAP[1]}')
    for i, r in enumerate(range(FEE[0], FEE[1] + 1)):
        put(ws, f'G{r}', FEES[i] if i < len(FEES) else None, 'i')
    units = d['units']
    for i, r in enumerate(range(UNIT[0], UNIT[1] + 1)):
        u = units[i] if i < len(units) else (None, None, None)
        put(ws, f'I{r}', u[0], 'i'); put(ws, f'J{r}', u[1], 'i'); put(ws, f'K{r}', u[2], 'i')
        put(ws, f'L{r}', None, 'im'); put(ws, f'M{r}', None, 'im'); put(ws, f'N{r}', None, 'i')
    dv_list(ws, '"供应商,客户,供应商+客户,其他"', f'J{UNIT[0]}:J{UNIT[1]}')
    dv_list(ws, '"线上,线下"', f'K{UNIT[0]}:K{UNIT[1]}')
    goods = d['goods']
    for i, r in enumerate(range(GOOD[0], GOOD[1] + 1)):
        g = goods[i] if i < len(goods) else (None,) * 5
        for L, v, s in zip('PQRST', g, ('i', 'i', 'i', 'in', 'im')):
            put(ws, f'{L}{r}', v, s)
        put(ws, f'U{r}', f'=IF(P{r}="","",ROUND(N(S{r})*N(T{r}),2))', 'am')
    ws.freeze_panes = 'A5'
    # 下拉名称（OFFSET 只用在这里）
    for nm, c0, r0, r1 in (('L_账户', 'A', *ACC), ('L_类别', 'A', *CAT), ('L_费用', 'G', *FEE), ('L_往来', 'I', *UNIT)):
        wb.defined_names[nm] = DefinedName(nm, attr_text=f"OFFSET('{BS}'!${c0}${r0},0,0,MAX(1,COUNTA('{BS}'!${c0}${r0}:${c0}${r1})),1)")


# ─────────────────────────── 录入本：资金台帐 ───────────────────────────
def s_cash(wb, d):
    ws = wb[CA]
    head(ws, 3, [('A', '日期', 11), ('B', '账户', 14), ('C', '收支类别', 11), ('D', '往来单位', 20), ('E', '费用项目', 11),
                 ('F', '摘要', 26), ('G', '收入', 13), ('H', '支出', 13), ('I', '对方账户\n(内部转账填)', 14), ('J', '备注', 20),
                 ('K', '本账户余额', 14), ('L', '问题提示', 24), ('M', '属性', 8), ('N', '年月', 8)],
         'A076 资金台帐（招商银行／微信／支付宝／拼多多店铺余额 混合录入）',
         '一行一笔，按日期先后往下记（余额按行顺序累计）。收入、支出只填一边。内部转账（拼多多提现到招行、微信转招行等）：账户填转出方、支出填金额、对方账户填转入方，一行就够。'
         '往来款选往来单位，费用支出选费用项目。灰色列自动算；右边隐藏了 6 列各账户累计数（别删）。拼多多店铺余额这行的余额＝期初＋台帐＋账务明细（不含提现）到当天。')
    r0, r1 = CR
    for L, s in zip('ABCDEFGHIJ', ('id', 'i', 'i', 'i', 'i', 'i', 'im', 'im', 'i', 'i')):
        col(ws, L, r0, r1, None, s)
    for k, L in enumerate(SLOTS):
        a = f"'{BS}'!$A${ACC[0] + k}"
        put(ws, f'{L}3', f"=N('{BS}'!$C${ACC[0] + k})", 'hh')
        col(ws, L, r0, r1, f'={L}{{p}}+IF($B{{r}}={a},N($G{{r}})-N($H{{r}}),0)+IF($I{{r}}={a},N($H{{r}})-N($G{{r}}),0)', 'am')
        ws.column_dimensions[L].hidden = True
    acc = f"'{BS}'!$A${ACC[0]}:$A${ACC[1]}"
    col(ws, 'K', r0, r1, f'=IF($B{{r}}="","",IFERROR(INDEX($O{{r}}:$T{{r}},MATCH($B{{r}},{acc},0))'
                         f'+IF($B{{r}}={PDD},SUMIFS({RG(P1, "K")},{RG(P1, "H")},"<"&(INT(N($A{{r}}))+1),{RG(P1, "L")},"<>提现"),0),""))', 'am')
    col(ws, 'L', r0, r1,
        '=IF(AND($A{r}="",$B{r}="",N($G{r})=0,N($H{r})=0),"",IF(NOT(ISNUMBER($A{r})),"日期没填或不是日期",IF($B{r}="","缺账户",'
        f'IF(ISNA(MATCH($B{{r}},{acc},0)),"账户不在基础资料",IF($C{{r}}="","缺收支类别",IF($M{{r}}="?","类别不在基础资料",'
        'IF(AND(N($G{r})<>0,N($H{r})<>0),"收入、支出只填一边",IF(N($G{r})+N($H{r})=0,"缺金额",'
        'IF(AND($M{r}="内部转账",$I{r}=""),"内部转账要填对方账户",IF(AND($I{r}<>"",$I{r}=$B{r}),"对方账户不能是本账户",'
        f'IF(AND($I{{r}}<>"",ISNA(MATCH($I{{r}},{acc},0))),"对方账户不在基础资料",IF(AND($M{{r}}="往来",$D{{r}}=""),"往来款要选往来单位",'
        'IF(AND($M{r}="支出",$E{r}=""),"支出最好选费用项目",IF(N($K{r})<0,"余额为负：期初没填或漏记收入",""))))))))))))))', 'a')
    col(ws, 'M', r0, r1, f"=IF($C{{r}}=\"\",\"\",IFERROR(INDEX('{BS}'!$B${CAT[0]}:$B${CAT[1]},MATCH($C{{r}},'{BS}'!$A${CAT[0]}:$A${CAT[1]},0)),\"?\"))", 'a')
    col(ws, 'N', r0, r1, '=IF(ISNUMBER($A{r}),YEAR($A{r})*100+MONTH($A{r}),"")', 'ai')
    dv_list(ws, 'L_账户', f'B{r0}:B{r1}'); dv_list(ws, 'L_账户', f'I{r0}:I{r1}')
    dv_list(ws, 'L_类别', f'C{r0}:C{r1}'); dv_list(ws, 'L_往来', f'D{r0}:D{r1}'); dv_list(ws, 'L_费用', f'E{r0}:E{r1}')
    for i, row in enumerate(d['cash']):
        for L, v in zip('ABCDEFGHIJ', row):
            if v is not None:
                ws[f'{L}{r0 + i}'].value = v
    ws.freeze_panes = 'C4'
    ws.auto_filter.ref = f'A3:N{r1}'


# ─────────────────────────── 录入本：采购入库、销售出库（管家婆导出原样列） ───────────────────────────
PUR_H = ['行号', '录单时间', '单据编号', '单位名称', '仓库', '经手人', '附加说明', '商品编号', '商品名称', '规格', '型号', '品牌',
         '序列号', '数量', '单位', '单价', '金额', '明细备注']
SAL_H = ['行号', '单位名称', '录单日期', '单据类型', '单据编号', '明细备注', '仓库', '商品名称', '序列号', '单据明细数量', '单位', '单价', '金额']
DATE_F = '=IF({c}{{r}}="","",IFERROR(INT(--{c}{{r}}),IFERROR(DATE(LEFT({c}{{r}},4),MID({c}{{r}},6,2),MID({c}{{r}},9,2)),"")))'


def _paste_cols(ws, hdrs, widths, row=3):
    cols = []
    for i, h in enumerate(hdrs):
        cols.append((CL(i + 1), h, widths.get(h, 11)))
    return cols


def s_pur(wb, d):
    ws = wb[PU]
    cols = _paste_cols(ws, PUR_H, {'行号': 6, '录单时间': 17, '单据编号': 18, '单位名称': 20, '商品名称': 30, '序列号': 18, '明细备注': 12, '附加说明': 8})
    cols += [('S', '日期', 11), ('T', '年月', 8), ('U', '数量(数)', 11), ('V', '金额(数)', 13), ('W', '问题提示', 18)]
    head(ws, 3, cols, 'A076 采购入库（直接粘管家婆「商品采购明细查询」）',
         '管家婆导出后，从明细第 1 行（表头下一行）复制到最后一行明细（合计行可以不要），粘到 A4。A～R 列跟管家婆导出一模一样；S 列起是灰色自动列，别粘过去。'
         '采购退货在管家婆导出里是负数，原样粘。新商品要先在基础资料加上，库存表才算得到。')
    r0, r1 = PR
    for i in range(18):
        col(ws, CL(i + 1), r0, r1, None, 'i')
    col(ws, 'S', r0, r1, DATE_F.format(c='B'), 'ad')
    col(ws, 'T', r0, r1, '=IF(S{r}="","",YEAR(S{r})*100+MONTH(S{r}))', 'ai')
    col(ws, 'U', r0, r1, '=IF(S{r}="",0,IFERROR(--N{r},0))', 'an')
    col(ws, 'V', r0, r1, '=IF(S{r}="",0,IFERROR(--Q{r},0))', 'am')
    col(ws, 'W', r0, r1, f"=IF(S{{r}}=\"\",\"\",IF(I{{r}}=\"\",\"缺商品名称\",IF(ISNA(MATCH(I{{r}},'{BS}'!$P${GOOD[0]}:$P${GOOD[1]},0)),\"商品不在基础资料\","
                         f"IF(ISNA(MATCH(D{{r}},'{BS}'!$I${UNIT[0]}:$I${UNIT[1]},0)),\"供应商不在基础资料\",IF(ABS(U{{r}}*IFERROR(--P{{r}},0)-V{{r}})>1,\"数量×单价≠金额\",\"\")))))", 'a')
    for i, row in enumerate(d['pur']):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=r0 + i, column=j + 1).value = v
    ws.freeze_panes = 'C4'
    ws.auto_filter.ref = f'A3:W{r1}'


def s_sal(wb, d):
    ws = wb[SA]
    cols = _paste_cols(ws, SAL_H, {'行号': 6, '单位名称': 22, '录单日期': 17, '单据类型': 14, '单据编号': 18, '商品名称': 30, '序列号': 16, '单据明细数量': 11})
    cols += [('N', '渠道(可不填)', 10), ('O', '渠道', 8), ('P', '日期', 11), ('Q', '年月', 8), ('R', '数量(数)', 11), ('S', '金额(数)', 13), ('T', '问题提示', 18)]
    head(ws, 3, cols, 'A076 销售出库（直接粘管家婆「销售明细查询-按商品统计」）',
         '管家婆导出后，从明细第 1 行复制到最后一行明细，粘到 A4（A～M 跟导出一样；导出里第 14 列是空的，粘过来落在 N 列也没关系）。'
         'N 列「渠道」可以不填：单位名称带「拼多多」或往来单位渠道是「线上」的算拼多多，其余算线下；要改就在 N 列选。O 列起灰色自动。')
    r0, r1 = PR
    for i in range(13):
        col(ws, CL(i + 1), r0, r1, None, 'i')
    col(ws, 'N', r0, r1, None, 'i')
    dv_list(ws, '"拼多多,线下"', f'N{r0}:N{r1}')
    col(ws, 'O', r0, r1, f"=IF(P{{r}}=\"\",\"\",IF(N{{r}}<>\"\",N{{r}},IF(ISNUMBER(FIND(\"拼多多\",B{{r}})),\"拼多多\","
                         f"IF(IFERROR(INDEX('{BS}'!$K${UNIT[0]}:$K${UNIT[1]},MATCH(B{{r}},'{BS}'!$I${UNIT[0]}:$I${UNIT[1]},0)),\"\")=\"线上\",\"拼多多\",\"线下\"))))", 'a')
    col(ws, 'P', r0, r1, DATE_F.format(c='C'), 'ad')
    col(ws, 'Q', r0, r1, '=IF(P{r}="","",YEAR(P{r})*100+MONTH(P{r}))', 'ai')
    col(ws, 'R', r0, r1, '=IF(P{r}="",0,IFERROR(--J{r},0))', 'an')
    col(ws, 'S', r0, r1, '=IF(P{r}="",0,IFERROR(--M{r},0))', 'am')
    col(ws, 'T', r0, r1, f"=IF(P{{r}}=\"\",\"\",IF(H{{r}}=\"\",\"缺商品名称\",IF(ISNA(MATCH(H{{r}},'{BS}'!$P${GOOD[0]}:$P${GOOD[1]},0)),\"商品不在基础资料\","
                         f"IF(ISNA(MATCH(B{{r}},'{BS}'!$I${UNIT[0]}:$I${UNIT[1]},0)),\"客户不在基础资料\",IF(AND(O{{r}}<>\"拼多多\",O{{r}}<>\"线下\"),\"渠道只能是拼多多或线下\",\"\")))))", 'a')
    for i, row in enumerate(d['sal']):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=r0 + i, column=j + 1).value = v
    ws.freeze_panes = 'C4'
    ws.auto_filter.ref = f'A3:T{r1}'


# ─────────────────────────── 录入本：拼多多三张粘贴区 ───────────────────────────
P1_H = ['商户订单号', '发生时间', '收入金额（+元）', '支出金额（-元）', '账务类型', '备注', '业务描述']
BILL_H = ['订单号', '账单类型', '活动类型', '商品ID', '商品名称', '商品售价(元)', '成本金额(元)', '结算总金额(元)', '批次号', '结算方式', '结算时间', '备注']


def s_p1(wb, d):
    ws = wb[P1]
    cols = [(CL(i + 1), h, w) for i, (h, w) in enumerate(zip(P1_H, (24, 17, 12, 12, 11, 40, 30)))]
    cols += [('H', '时间', 16), ('I', '年月', 8), ('J', '#总计', 12), ('K', '净额', 12), ('L', '归到', 14), ('M', '批次号', 19),
             ('N', '批次序号', 7), ('O', '百亿扣费序号', 7)]
    head(ws, 3, cols, 'A076 拼多多店铺账务明细（直接粘拼多多导出）',
         '拼多多后台「账务明细」导出后，从「商户订单号」表头下一行复制到最后（#支出合计、#收入合计、#总计 那几行一起粘，用来核对），粘到 A4。'
         '下个月接着粘在上个月下面（中间空行没关系）。逐单核对技术服务费要跨月，所以最好把最近几个月都留着。H 列起灰色自动，别粘过去。')
    r0, r1 = XR
    for i in range(7):
        col(ws, CL(i + 1), r0, r1, None, 'i')
    col(ws, 'H', r0, r1, '=IF(ISNUMBER(B{r}),B{r},IF(B{r}="","",IFERROR(--B{r},"")))', 'adt')
    col(ws, 'I', r0, r1, '=IF(H{r}<>"",YEAR(H{r})*100+MONTH(H{r}),N(I{p}))', 'ai')
    col(ws, 'J', r0, r1, '=IF(LEFT(A{r},3)="#总计",IFERROR(--SUBSTITUTE(MID(A{r},5,99),"元",""),""),"")', 'am')
    col(ws, 'K', r0, r1, '=IF(H{r}="",0,IFERROR(--C{r},0)+IFERROR(--D{r},0))', 'am')
    col(ws, 'L', r0, r1, f"=IF(H{{r}}=\"\",\"\",IFERROR(INDEX('{BS}'!$B${MAP[0]}:$B${MAP[1]},MATCH(LEFT(G{{r}},7),'{BS}'!$A${MAP[0]}:$A${MAP[1]},0)),\"其他未归类\"))", 'a')
    col(ws, 'M', r0, r1, '=IF(H{r}="","",IFERROR(MID(F{r},FIND("BO",F{r}),17),""))', 'a')
    col(ws, 'N', r0, r1, '=N(N{p})+(M{r}<>"")', 'ai')
    col(ws, 'O', r0, r1, '=N(O{p})+AND(L{r}="百亿补贴技术服务费",K{r}<0)', 'ai')
    for i, row in enumerate(d['p1']):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=r0 + i, column=j + 1).value = v
    ws.freeze_panes = 'C4'
    ws.auto_filter.ref = f'A3:O{r1}'


def _bill(ws, title, tip, extra, rows):
    cols = [(CL(i + 1), h, w) for i, (h, w) in enumerate(zip(BILL_H, (24, 10, 12, 13, 30, 10, 10, 11, 19, 9, 17, 16)))]
    cols += extra
    head(ws, 3, cols, title, tip)
    r0, r1 = XR
    for i in range(12):
        col(ws, CL(i + 1), r0, r1, None, 'i')
    for i, row in enumerate(rows):
        for j, v in enumerate(row):
            if v is not None:
                ws.cell(row=r0 + i, column=j + 1).value = v
    ws.freeze_panes = 'B4'


def s_p2(wb, d):
    ws = wb[P2]
    _bill(ws, 'A076 拼多多营销活动结算对账单（直接粘）',
          '拼多多后台导出的「营销活动结算」对账单，从表头下一行起粘到 A4，多次导出接着往下粘（同一单别重复粘）。'
          '公式：应结＝成本金额（报名价）－商品售价（活动价）＋备注里的「补贴优惠券X元」，跟「结算总金额」比；批次合计到「拼多多对账」表跟账务明细的补贴汇入比。',
          [('M', '结算额(数)', 11), ('N', '补贴券', 9), ('O', '应结', 11), ('P', '差额', 10), ('Q', '结论', 8), ('R', '批次序号', 7), ('S', '不一致序号', 7)],
          d['p2'])
    r0, r1 = XR
    col(ws, 'M', r0, r1, '=IF(A{r}="","",IFERROR(--H{r},0))', 'am')
    col(ws, 'N', r0, r1, '=IF(A{r}="","",IFERROR(--MID(L{r},FIND("补贴优惠券",L{r})+5,FIND("元",L{r},FIND("补贴优惠券",L{r}))-FIND("补贴优惠券",L{r})-5),0))', 'am')
    col(ws, 'O', r0, r1, '=IF(A{r}="","",IFERROR(--G{r},0)-IFERROR(--F{r},0)+N{r})', 'am')
    col(ws, 'P', r0, r1, '=IF(A{r}="","",ROUND(M{r}-O{r},2))', 'am')
    col(ws, 'Q', r0, r1, '=IF(A{r}="","",IF(ABS(P{r})<0.005,"对","不一致"))', 'a')
    col(ws, 'R', r0, r1, f'=N(R{{p}})+IF(I{{r}}="",0,IF(MATCH(I{{r}},$I${r0}:$I${r1},0)=ROW()-{r0 - 1},1,0))', 'ai')
    col(ws, 'S', r0, r1, '=N(S{p})+(Q{r}="不一致")', 'ai')
    red_if(ws, f'Q{r0}:Q{r1}', f'Q{r0}="不一致"')
    ws.auto_filter.ref = f'A3:S{r1}'


def s_p3(wb, d):
    ws = wb[P3]
    _bill(ws, 'A076 拼多多技术服务费返还账单（直接粘）',
          '拼多多后台导出的「服务费结算／技术服务费返还」账单，从表头下一行起粘到 A4，多批接着往下粘。订单号带「-t01」的是返了又扣回（负数）。'
          '每单应返＝百亿补贴技术服务费（订单价×1.4%）；逐单跟账务明细里扣的费对，在「技术服务费核对」表看。',
          [('M', '金额(数)', 10), ('N', '原订单号', 24), ('O', '返还/扣回', 9), ('P', '账务明细里有这单', 9), ('Q', '返还÷1.4%\n=订单价', 11),
           ('R', '批次序号', 7), ('S', '扣回序号', 7)],
          d['p3'])
    r0, r1 = XR
    col(ws, 'M', r0, r1, '=IF(A{r}="","",IFERROR(--H{r},0))', 'am')
    col(ws, 'N', r0, r1, '=IF(A{r}="","",IF(ISNUMBER(FIND("-t",A{r})),LEFT(A{r},FIND("-t",A{r})-1),TRIM(A{r})))', 'a')
    col(ws, 'O', r0, r1, '=IF(A{r}="","",IF(OR(N{r}<>TRIM(A{r}),N(M{r})<0),"扣回","返还"))', 'a')
    col(ws, 'P', r0, r1, f'=IF(A{{r}}="","",IF(ISNUMBER(MATCH(N{{r}},{RG(P1, "A")},0)),"有","没有"))', 'a')
    col(ws, 'Q', r0, r1, '=IF(A{r}="","",ROUND(ABS(N(M{r}))/0.014,0))', 'am')
    col(ws, 'R', r0, r1, f'=N(R{{p}})+IF(I{{r}}="",0,IF(MATCH(I{{r}},$I${r0}:$I${r1},0)=ROW()-{r0 - 1},1,0))', 'ai')
    col(ws, 'S', r0, r1, '=N(S{p})+(O{r}="扣回")', 'ai')
    red_if(ws, f'O{r0}:O{r1}', f'O{r0}="扣回"')
    ws.auto_filter.ref = f'A3:S{r1}'


# ─────────────────────────── 报表本 ───────────────────────────
YM = f"'{HOME}'!$B$4*100"          # 年份*100（再加月份）


def month_hdr(ws, row, c0):
    for m in range(1, 13):
        put(ws, f'{CL(CI(c0) + m - 1)}{row}', m, 'h').number_format = '0"月"'
        ws.column_dimensions[CL(CI(c0) + m - 1)].width = 11


def s_rec(wb, d):
    ws = wb[REC]
    head(ws, 5, [('A', '账务类型', 18), ('B', '利润表归类', 16)] + [(CL(i), '', 11) for i in range(3, 16)],
         'A076 拼多多对账（店铺余额·批次·营销活动）',
         '怎么核才准：① 先看「一」每月净额跟账务明细最后一行「#总计」对上没有；② 补贴汇入、技术服务费返还在账务明细里只有一笔总数，备注里有批次号——拿批次号去对账单里找同批次的明细，合计要等于到账金额（「二」）；'
         '③ 营销活动逐单：应结＝报名价－活动价＋补贴券（「三」）；④ 技术服务费逐单在「技术服务费核对」表。对账单是按结算时间出的，账务明细是按到账月份出的，月份不一样就会「对不上」，要按批次号对。', last='O')
    put(ws, 'A4', '一、按月店铺余额（年份见首页；单位：元，收入正、支出负）', 'sec')
    put(ws, 'A5', '账务类型', 'h'); put(ws, 'B5', '利润表归类', 'h'); month_hdr(ws, 5, 'C'); put(ws, 'O5', '全年合计', 'h')
    P1K, P1L, P1I = RG(P1, 'K'), RG(P1, 'L'), RG(P1, 'I')
    rr = list(range(6, 6 + MAP[1] - MAP[0] + 1))          # 6..20
    for i, r in enumerate(rr):
        m = MAP[0] + i
        put(ws, f'A{r}', f"=IF('{BS}'!$B${m}=\"\",\"\",'{BS}'!$B${m})", 'a')
        put(ws, f'B{r}', f"=IF(A{r}=\"\",\"\",'{BS}'!$C${m})", 'a')
    r_oth = rr[-1] + 1
    put(ws, f'A{r_oth}', '其他未归类', 'a'); put(ws, f'B{r_oth}', '其他费用', 'a')
    for r in rr + [r_oth]:
        for m in range(1, 13):
            L = CL(2 + m)
            put(ws, f'{L}{r}', f'=IF($A{r}="",0,SUMIFS({P1K},{P1L},$A{r},{P1I},{YM}+{L}$5))', 'am')
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})', 'am')
    R = dict(net=r_oth + 1, tot=r_oth + 2, diff=r_oth + 3, beg=r_oth + 4, end=r_oth + 5, cash=r_oth + 6, cdiff=r_oth + 7)
    labels = dict(net='本月净额（上面合计）', tot='账务明细「#总计」', diff='差异（应为 0）', beg='月初店铺余额', end='月末店铺余额',
                  cash='资金台帐记的提现', cdiff='提现差异（台帐－账务明细）')
    CAB, CAG, CAH, CAI, CAM, CAN = (RG(CA, c) for c in 'BGHIMN')
    for k, r in R.items():
        put(ws, f'A{r}', labels[k], 'tot'); put(ws, f'B{r}', None, 'tot')
        for m in range(1, 13):
            L = CL(2 + m)
            ym = f'{YM}+{L}$5'
            f = {'net': f'=SUM({L}{rr[0]}:{L}{r_oth})',
                 'tot': f'=IF(COUNTIFS({RG(P1, "J")},">-9E+99",{P1I},{ym})=0,"",SUMIFS({RG(P1, "J")},{P1I},{ym}))',
                 'diff': f'=IF({L}{R["tot"]}="","没粘#总计",ROUND({L}{R["net"]}-{L}{R["tot"]},2))',
                 'beg': f"=N('{BS}'!$C$8)" if m == 1 else f'={CL(1 + m)}{R["end"]}',
                 'end': f'={L}{R["beg"]}+{L}{R["net"]}',
                 'cash': f'=SUMIFS({CAH},{CAB},{PDD},{CAM},"内部转账",{CAN},{ym})-SUMIFS({CAG},{CAB},{PDD},{CAM},"内部转账",{CAN},{ym})'
                         f'+SUMIFS({CAG},{CAI},{PDD},{CAM},"内部转账",{CAN},{ym})-SUMIFS({CAH},{CAI},{PDD},{CAM},"内部转账",{CAN},{ym})',
                 'cdiff': f'=ROUND({L}{R["cash"]}+SUMIFS({P1K},{P1L},"提现",{P1I},{ym}),2)'}[k]
            put(ws, f'{L}{r}', f, 'totm')
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})' if k in ('net', 'cash', 'cdiff') else None, 'totm')
    red_if(ws, f'C{R["diff"]}:N{R["diff"]}', f'AND(ISNUMBER(C{R["diff"]}),C{R["diff"]}<>0)')
    red_if(ws, f'C{R["cdiff"]}:N{R["cdiff"]}', f'C{R["cdiff"]}<>0')

    # 二 A：账务明细里的到账批次
    a0 = R['cdiff'] + 2
    put(ws, f'A{a0}', '二、批次核对　A. 账务明细里到账的批次（补贴汇入／技术服务费返还）→ 对账单里同一批次号的合计', 'sec')
    put(ws, f'P{a0}', f"={X(P1, f'$N${XR[1]}')}", 'ai'); put(ws, f'Q{a0}', '← 批数', 'lb')
    h = a0 + 1
    for L, t in zip('ABCDEFGHIJP', ('序', '到账时间', '类型', '批次号', '到账金额', '营销活动对账单', '服务费返还账单', '对账单笔数', '差额', '结论', '(位置)')):
        put(ws, f'{L}{h}', t, 'h')
    A_rows = range(h + 1, h + 61)
    for i, r in enumerate(A_rows, 1):
        put(ws, f'A{r}', i, 'ai')
        put(ws, f'P{r}', f'=IF(A{r}>$P${a0},"",MATCH(A{r},{RG(P1, "N")},0))', 'hh')
        g = lambda c: f'=IF($P{r}="","",INDEX({RG(P1, c)},$P{r}))'
        put(ws, f'B{r}', g('H'), 'adt'); put(ws, f'C{r}', g('L'), 'a'); put(ws, f'D{r}', g('M'), 'a'); put(ws, f'E{r}', g('K'), 'am')
        put(ws, f'F{r}', f'=IF($P{r}="","",SUMIFS({RG(P2, "M")},{RG(P2, "I")},$D{r}))', 'am')
        put(ws, f'G{r}', f'=IF($P{r}="","",SUMIFS({RG(P3, "M")},{RG(P3, "I")},$D{r}))', 'am')
        put(ws, f'H{r}', f'=IF($P{r}="","",COUNTIFS({RG(P2, "I")},$D{r})+COUNTIFS({RG(P3, "I")},$D{r}))', 'ai')
        put(ws, f'I{r}', f'=IF($P{r}="","",ROUND(E{r}-F{r}-G{r},2))', 'am')
        put(ws, f'J{r}', f'=IF($P{r}="","",IF(H{r}=0,"对账单里没有这批明细",IF(ABS(I{r})<0.01,"对上","差 "&TEXT(I{r},"0.00"))))', 'a')
    A_last = h + 60
    red_if(ws, f'J{h + 1}:J{A_last}', f'AND(J{h + 1}<>"",J{h + 1}<>"对上")')
    green_if(ws, f'J{h + 1}:J{A_last}', f'J{h + 1}="对上"')

    # 二 B：对账单里的批次
    b0 = A_last + 2
    put(ws, f'A{b0}', 'B. 对账单里的批次 → 账务明细里到账了没有', 'sec')
    put(ws, f'P{b0}', f"={X(P2, f'$R${XR[1]}')}", 'ai'); put(ws, f'Q{b0}', f"={X(P3, f'$R${XR[1]}')}", 'ai')
    put(ws, f'R{b0}', '← 营销活动批数、服务费返还批数', 'lb')
    h = b0 + 1
    for L, t in zip('ABCDEFGHIP', ('序', '来源', '批次号', '结算时间', '笔数', '对账单合计', '账务明细到账', '差额', '结论', '(位置)')):
        put(ws, f'{L}{h}', t, 'h')
    B_rows = range(h + 1, h + 41)
    for i, r in enumerate(B_rows, 1):
        put(ws, f'A{r}', i, 'ai')
        put(ws, f'B{r}', f'=IF(A{r}<=$P${b0},"营销活动",IF(A{r}<=$P${b0}+$Q${b0},"服务费返还",""))', 'a')
        put(ws, f'P{r}', f'=IF(B{r}="营销活动",MATCH(A{r},{RG(P2, "R")},0),IF(B{r}="服务费返还",MATCH(A{r}-$P${b0},{RG(P3, "R")},0),""))', 'hh')
        g = lambda c: f'=IF($P{r}="","",IF($B{r}="营销活动",INDEX({RG(P2, c)},$P{r}),INDEX({RG(P3, c)},$P{r})))'
        put(ws, f'C{r}', g('I'), 'a'); put(ws, f'D{r}', g('K'), 'a')
        put(ws, f'E{r}', f'=IF($P{r}="","",COUNTIFS({RG(P2, "I")},$C{r})+COUNTIFS({RG(P3, "I")},$C{r}))', 'ai')
        put(ws, f'F{r}', f'=IF($P{r}="","",SUMIFS({RG(P2, "M")},{RG(P2, "I")},$C{r})+SUMIFS({RG(P3, "M")},{RG(P3, "I")},$C{r}))', 'am')
        put(ws, f'G{r}', f'=IF($P{r}="","",SUMIFS({P1K},{RG(P1, "M")},$C{r}))', 'am')
        put(ws, f'H{r}', f'=IF($P{r}="","",ROUND(G{r}-F{r},2))', 'am')
        put(ws, f'I{r}', f'=IF($P{r}="","",IF(G{r}=0,"账务明细里还没到账（看结算时间，粘那个月的账务明细）",IF(ABS(H{r})<0.01,"对上","差 "&TEXT(H{r},"0.00"))))', 'a')
    B_last = h + 40
    red_if(ws, f'I{h + 1}:I{B_last}', f'AND(I{h + 1}<>"",I{h + 1}<>"对上")')
    green_if(ws, f'I{h + 1}:I{B_last}', f'I{h + 1}="对上"')

    # 三：营销活动逐单
    c0 = B_last + 2
    put(ws, f'A{c0}', '三、营销活动逐单核对：应结＝成本金额（报名价）－商品售价（活动价）＋备注里的补贴优惠券，跟结算总金额比', 'sec')
    s = c0 + 1
    for L, t, f, st in (('A', '对账单笔数', f'=COUNT({RG(P2, "M")})', 'ai'), ('C', '应结合计', f'=SUM({RG(P2, "O")})', 'am'),
                        ('E', '结算合计', f'=SUM({RG(P2, "M")})', 'am'), ('G', '补贴券合计', f'=SUM({RG(P2, "N")})', 'am'),
                        ('I', '不一致笔数', f'=COUNTIFS({RG(P2, "Q")},"不一致")', 'ai')):
        put(ws, f'{L}{s}', t, 'lb'); put(ws, f'{L}{s + 1}', f, st)
    h = s + 3
    put(ws, f'A{h - 1}', '不一致的单（最多列 20 笔，全部的去「粘贴_营销活动」按「结论」筛选）', 'lb')
    for L, t in zip('ABCDEFGHIJP', ('序', '订单号', '商品名称', '活动价', '报名价', '补贴券', '应结', '结算总金额', '差额', '批次号', '(位置)')):
        put(ws, f'{L}{h}', t, 'h')
    for i, r in enumerate(range(h + 1, h + 21), 1):
        put(ws, f'A{r}', i, 'ai')
        put(ws, f'P{r}', f'=IFERROR(MATCH(A{r},{RG(P2, "S")},0),"")', 'hh')
        for L, c, st in (('B', 'A', 'a'), ('C', 'E', 'a'), ('D', 'F', 'am'), ('E', 'G', 'am'), ('F', 'N', 'am'), ('G', 'O', 'am'),
                         ('H', 'M', 'am'), ('I', 'P', 'am'), ('J', 'I', 'a')):
            put(ws, f'{L}{r}', f'=IF($P{r}="","",INDEX({RG(P2, c)},$P{r}))', st)
    for L, w in (('B', 16), ('D', 19), ('I', 14), ('J', 24), ('P', 7), ('Q', 7)):
        ws.column_dimensions[L].width = w
    ws.freeze_panes = 'C6'
    return dict(R=R, rr=rr, r_oth=r_oth, a0=a0, A=(a0 + 2, A_last), b0=b0, B=(b0 + 2, B_last), c0=c0 + 2)


def s_tf(wb, d):
    ws = wb[TF]
    cols = [('A', '序', 6), ('B', '订单号', 24), ('C', '扣费时间', 16), ('D', '订单价', 11), ('E', '应返(1.4%)', 11), ('F', '实扣百亿费', 11),
            ('G', '已返还', 11), ('H', '扣回', 10), ('I', '待返', 11), ('J', '状态', 30), ('K', '距账务明细\n最后一天(天)', 11), ('L', '(位置)', 7)]
    head(ws, 28, cols, 'A076 技术服务费核对（百亿补贴技术服务费 1.4% 返还，逐单）',
         '百亿补贴的单，确认收货时账务明细扣「基础技术服务费 0.6%」＋「百亿补贴技术服务费 1.4%」；1.4% 这部分平台过后按批次返还（返还账单一单一行，账务明细里只有一笔批次总额）。'
         '下面「三」逐单看：扣了多少、返了多少、有没有返了又扣回（订单号带 -t01）。返还一般晚一个月左右，所以要把扣费那个月和返还那个月的账务明细都粘进录入本。')
    P3M, P3N, P3O, P3P = (RG(P3, c) for c in 'MNOP')
    put(ws, 'A4', '一、汇总', 'sec')
    left = [('返还账单：返还笔数', f'=COUNTIFS({P3O},"返还")', 'ai'), ('返还合计', f'=SUMIFS({P3M},{P3O},"返还")', 'am'),
            ('扣回笔数（-t01）', f'=COUNTIFS({P3O},"扣回")', 'ai'), ('扣回合计', f'=SUMIFS({P3M},{P3O},"扣回")', 'am'),
            ('净返还', '=C6+C8', 'am'),
            ('账务明细里找不到这单（笔）', f'=COUNTIFS({P3P},"没有",{P3O},"返还")', 'ai'),
            ('　金额（扣费在更早月份）', f'=SUMIFS({P3M},{P3P},"没有",{P3O},"返还")', 'am')]
    rng = 'J29:J428'
    right = [('账务明细：扣了百亿费的订单数', f"={X(P1, f'$O${XR[1]}')}", 'ai'), ('实扣百亿费合计', '=SUM(F29:F428)', 'am'),
             ('订单已退款、不用返（单）', f'=COUNTIF({rng},"订单已退款*")', 'ai'),
             ('已返还（单）', f'=COUNTIF({rng},"已返还")', 'ai'), ('已返还金额', f'=SUMIFS(G29:G428,{rng},"已返还")', 'am'),
             ('扣了还没返（单）', f'=COUNTIF({rng},"扣了还没返")', 'ai'), ('扣了还没返金额', f'=SUMIFS(I29:I428,{rng},"扣了还没返")', 'am'),
             ('少返／返了又扣回（单）', f'=COUNTIF({rng},"少返*")+COUNTIF({rng},"返了又扣回*")', 'ai')]
    for i, (t, f, st) in enumerate(left):
        put(ws, f'B{5 + i}', t, 'lb'); put(ws, f'C{5 + i}', f, st)
    for i, (t, f, st) in enumerate(right):
        put(ws, f'E{5 + i}', t, 'lb'); put(ws, f'F{5 + i}', None, 'lb'); put(ws, f'G{5 + i}', None, 'lb')
        ws.merge_cells(f'E{5 + i}:G{5 + i}'); put(ws, f'H{5 + i}', f, st)
    put(ws, 'J5', '账务明细最后一天', 'lb'); put(ws, 'J6', f'=MAX({RG(P1, "H")})', 'adt')
    put(ws, 'A15', '二、返了又扣回（返还账单里订单号带 -t01 的负数行）', 'sec')
    for L, t in zip('ABCDEFL', ('序', '账单订单号', '原订单号', '金额', '批次号', '结算时间', '(位置)')):
        put(ws, f'{L}16', t, 'h')
    for i, r in enumerate(range(17, 25), 1):
        put(ws, f'A{r}', i, 'ai')
        put(ws, f'L{r}', f'=IFERROR(MATCH(A{r},{RG(P3, "S")},0),"")', 'hh')
        for L, c, st in (('B', 'A', 'a'), ('C', 'N', 'a'), ('D', 'M', 'am'), ('E', 'I', 'a'), ('F', 'K', 'a')):
            put(ws, f'{L}{r}', f'=IF($L{r}="","",INDEX({RG(P3, c)},$L{r}))', st)
    put(ws, 'A27', '三、逐单：账务明细里扣了「百亿补贴技术服务费」的订单（最多 400 单）', 'sec')
    P1A, P1K, P1L = RG(P1, 'A'), RG(P1, 'K'), RG(P1, 'L')
    for i, r in enumerate(range(29, 429), 1):
        put(ws, f'A{r}', i, 'ai')
        put(ws, f'L{r}', f"=IF(A{r}>$H$5,\"\",MATCH(A{r},{RG(P1, 'O')},0))", 'hh')
        put(ws, f'B{r}', f'=IF($L{r}="","",INDEX({P1A},$L{r}))', 'a')
        put(ws, f'C{r}', f'=IF($L{r}="","",INDEX({RG(P1, "H")},$L{r}))', 'adt')
        put(ws, f'D{r}', f'=IF($L{r}="","",SUMIFS({P1K},{P1A},$B{r},{P1L},"订单收入")+SUMIFS({P1K},{P1A},$B{r},{P1L},"优惠券结算"))', 'am')
        put(ws, f'E{r}', f'=IF($L{r}="","",ROUND(D{r}*0.014,2))', 'am')
        put(ws, f'F{r}', f'=IF($L{r}="","",-SUMIFS({P1K},{P1A},$B{r},{P1L},"百亿补贴技术服务费"))', 'am')
        put(ws, f'G{r}', f'=IF($L{r}="","",SUMIFS({P3M},{P3N},$B{r},{P3O},"返还"))', 'am')
        put(ws, f'H{r}', f'=IF($L{r}="","",SUMIFS({P3M},{P3N},$B{r},{P3O},"扣回"))', 'am')
        put(ws, f'I{r}', f'=IF($L{r}="","",ROUND(F{r}-G{r}-H{r},2))', 'am')
        put(ws, f'J{r}', f'=IF($L{r}="","",IF(F{r}<0.005,"订单已退款，百亿费已原路退回，不用返",IF(G{r}+H{r}>=F{r}-0.01,"已返还",'
                         f'IF(H{r}<0,"返了又扣回(-t01)，问客服",IF(G{r}>0,"少返 "&TEXT(I{r},"0.00"),"扣了还没返")))))', 'a')
        put(ws, f'K{r}', f'=IF($L{r}="","",INT($J$6)-INT(C{r}))', 'ai')
    red_if(ws, 'J29:J428', 'OR(J29="扣了还没返",LEFT(J29,2)="少返",LEFT(J29,2)="返了")')
    green_if(ws, 'J29:J428', 'J29="已返还"')
    ws.freeze_panes = 'C29'
    ws.auto_filter.ref = 'A28:K428'


def s_inv(wb, d):
    ws = wb[INV]
    cols = [('A', '商品名称', 40), ('B', '品牌', 8), ('C', '期初数量', 9), ('D', '期初金额', 13), ('E', '采购数量', 9), ('F', '采购金额', 13),
            ('G', '销售数量', 9), ('H', '加权单价', 11), ('I', '销售成本', 13), ('J', '结存数量', 9), ('K', '结存金额', 13), ('L', '提示', 18)]
    head(ws, 4, cols + [(CL(13 + i), '', 9) for i in range(12)], 'A076 库存（期初＋采购－销售＝结存）',
         '成本算法：全期一次加权平均——每个商品 加权单价＝（期初金额＋采购金额）÷（期初数量＋采购数量），销售成本＝销售数量×加权单价。'
         '期初取的是上传的管家婆库存表（你说不太准），盘点后到录入本「基础资料」改期初数量、单价。右边 12 列是年份内各月销售数量（利润表算各月成本用）。'
         '采购、销售都是全部明细（不按年份截）；建议一年一套录入本，年底把结存抄成下一年期初。', last='X')
    month_hdr(ws, 4, 'M')
    g0, g1 = GOOD
    r0 = 6
    PUI, PUU, PUV = RG(PU, 'I'), RG(PU, 'U'), RG(PU, 'V')
    SAH, SAR, SAQ = RG(SA, 'H'), RG(SA, 'R'), RG(SA, 'Q')
    for i in range(g1 - g0 + 1):
        r, b = r0 + i, g0 + i
        put(ws, f'A{r}', f"=IF('{BS}'!$P${b}=\"\",\"\",'{BS}'!$P${b})", 'a')
        put(ws, f'B{r}', f"=IF(A{r}=\"\",\"\",'{BS}'!$Q${b}&\"\")", 'a')
        put(ws, f'C{r}', f"=IF(A{r}=\"\",0,N('{BS}'!$S${b}))", 'an')
        put(ws, f'D{r}', f"=IF(A{r}=\"\",0,N('{BS}'!$U${b}))", 'am')
        put(ws, f'E{r}', f'=IF(A{r}="",0,SUMIFS({PUU},{PUI},$A{r}))', 'an')
        put(ws, f'F{r}', f'=IF(A{r}="",0,SUMIFS({PUV},{PUI},$A{r}))', 'am')
        put(ws, f'G{r}', f'=IF(A{r}="",0,SUMIFS({SAR},{SAH},$A{r}))', 'an')
        put(ws, f'H{r}', f'=IF(C{r}+E{r}>0,ROUND((D{r}+F{r})/(C{r}+E{r}),4),0)', 'am')
        put(ws, f'I{r}', f'=ROUND(G{r}*H{r},2)', 'am')
        put(ws, f'J{r}', f'=C{r}+E{r}-G{r}', 'an')
        put(ws, f'K{r}', f'=ROUND(D{r}+F{r}-I{r},2)', 'am')
        put(ws, f'L{r}', f'=IF(A{r}="","",IF(J{r}<0,"负库存：漏记采购或期初不准",IF(AND(G{r}>0,C{r}+E{r}=0),"没进货就卖了","")))', 'a')
        for m in range(1, 13):
            L = CL(12 + m)
            put(ws, f'{L}{r}', f'=IF($A{r}="",0,SUMIFS({SAR},{SAH},$A{r},{SAQ},{YM}+{L}$4))', 'an')
    last = r0 + g1 - g0
    put(ws, 'A5', '合计', 'tot'); put(ws, 'B5', None, 'tot')
    for L in 'CDEFGIJK':
        put(ws, f'{L}5', f'=SUM({L}{r0}:{L}{last})', 'totm')
    for L in ['H', 'L'] + [CL(12 + m) for m in range(1, 13)]:
        put(ws, f'{L}5', f'=SUM({L}{r0}:{L}{last})' if L not in ('H', 'L') else None, 'totm')
    put(ws, 'L5', f'=IF(ROUND(SUM({SAR})-G5,2)<>0,"有 "&ROUND(SUM({SAR})-G5,2)&" 个销售数量的商品不在基础资料","")'
                  f'&IF(ROUND(SUM({PUU})-E5,2)<>0,"；采购有 "&ROUND(SUM({PUU})-E5,2)&" 个不在基础资料","")', 'warn')
    red_if(ws, f'L{r0}:L{last}', f'L{r0}<>""')
    ws.freeze_panes = 'B6'
    ws.auto_filter.ref = f'A4:L{last}'


def s_ar(wb, d):
    ws = wb[AR]
    cols = [('A', '往来单位', 30), ('B', '类型', 11), ('C', '期初应付', 12), ('D', '采购额', 14), ('E', '已付款', 14), ('F', '期初应收', 12),
            ('G', '销售额', 14), ('H', '已收款', 14), ('I', '应付余额\n(我们欠对方)', 14), ('J', '应收余额\n(对方欠我们)', 14), ('K', '说明', 16)]
    head(ws, 4, cols, 'A076 往来（供应商应付／客户应收）',
         '采购额、销售额来自录入本的采购入库、销售出库（全部明细）；已付款、已收款来自资金台帐里「属性＝往来」且选了这个往来单位的行（付出去记支出、收进来记收入）。'
         '应付＝期初应付＋采购－付款；应收＝期初应收＋销售－收款；既是供应商又是客户的两边轧差。拼多多买家的货款走平台，看「拼多多对账」，这里不算应收。')
    u0, u1 = UNIT
    CAD, CAG, CAH, CAM = RG(CA, 'D'), RG(CA, 'G'), RG(CA, 'H'), RG(CA, 'M')
    r0 = 6
    for i in range(u1 - u0 + 1):
        r, b = r0 + i, u0 + i
        put(ws, f'A{r}', f"=IF('{BS}'!$I${b}=\"\",\"\",'{BS}'!$I${b})", 'a')
        put(ws, f'B{r}', f"=IF(A{r}=\"\",\"\",'{BS}'!$J${b}&IF('{BS}'!$K${b}=\"\",\"\",\"·\"&'{BS}'!$K${b}))", 'a')
        put(ws, f'C{r}', f"=IF(A{r}=\"\",0,N('{BS}'!$L${b}))", 'am')
        put(ws, f'D{r}', f'=IF(A{r}="",0,SUMIFS({RG(PU, "V")},{RG(PU, "D")},$A{r}))', 'am')
        put(ws, f'E{r}', f'=IF(A{r}="",0,SUMIFS({CAH},{CAD},$A{r},{CAM},"往来"))', 'am')
        put(ws, f'F{r}', f"=IF(A{r}=\"\",0,N('{BS}'!$M${b}))", 'am')
        put(ws, f'G{r}', f'=IF(OR(A{r}="",ISNUMBER(FIND("线上",B{r}))),0,SUMIFS({RG(SA, "S")},{RG(SA, "B")},$A{r}))', 'am')
        put(ws, f'H{r}', f'=IF(A{r}="",0,SUMIFS({CAG},{CAD},$A{r},{CAM},"往来"))', 'am')
        put(ws, f'I{r}', f'=MAX(0,ROUND((C{r}+D{r}-E{r})-(F{r}+G{r}-H{r}),2))', 'am')
        put(ws, f'J{r}', f'=MAX(0,ROUND((F{r}+G{r}-H{r})-(C{r}+D{r}-E{r}),2))', 'am')
        put(ws, f'K{r}', f'=IF(A{r}="","",IF(ISNUMBER(FIND("线上",B{r})),"平台结算，看拼多多对账",IF(I{r}>0,"我们欠对方",IF(J{r}>0,"对方欠我们",""))))', 'a')
    last = r0 + u1 - u0
    put(ws, 'A5', '合计', 'tot'); put(ws, 'B5', None, 'tot'); put(ws, 'K5', None, 'tot')
    for L in 'CDEFGHIJ':
        put(ws, f'{L}5', f'=SUM({L}{r0}:{L}{last})', 'totm')
    ws.freeze_panes = 'B6'
    ws.auto_filter.ref = f'A4:K{last}'


def s_exp(wb, d):
    ws = wb[EXP]
    head(ws, 4, [('A', '费用项目', 18)] + [(CL(i), '', 11) for i in range(2, 15)], 'A076 费用统计（费用项目 × 月，年份见首页）',
         '来自资金台帐：选了费用项目的行，支出－收入（退回的费用记收入会冲减）。最下面跟「属性＝支出」的合计比，差额＝没选费用项目、或费用项目不在基础资料的支出。'
         '拼多多扣的技术服务费、体验计划、售后补偿不在这里，在利润表「平台费用」。', last='N')
    month_hdr(ws, 4, 'B'); put(ws, 'N4', '全年合计', 'h')
    CAE, CAG, CAH, CAM, CAN = (RG(CA, c) for c in 'EGHMN')
    f0, f1 = FEE
    r0 = 5
    for i in range(f1 - f0 + 1):
        r, b = r0 + i, f0 + i
        put(ws, f'A{r}', f"=IF('{BS}'!$G${b}=\"\",\"\",'{BS}'!$G${b})", 'a')
        for m in range(1, 13):
            L = CL(1 + m)
            put(ws, f'{L}{r}', f'=IF($A{r}="",0,SUMIFS({CAH},{CAE},$A{r},{CAN},{YM}+{L}$4)-SUMIFS({CAG},{CAE},$A{r},{CAN},{YM}+{L}$4))', 'am')
        put(ws, f'N{r}', f'=SUM(B{r}:M{r})', 'am')
    last = r0 + f1 - f0
    t1, t2, t3 = last + 1, last + 2, last + 3
    put(ws, f'A{t1}', '按项目合计', 'tot'); put(ws, f'A{t2}', '台帐「属性＝支出」合计', 'tot'); put(ws, f'A{t3}', '差额（没选项目的支出）', 'tot')
    for m in range(1, 14):
        L = CL(1 + m)
        put(ws, f'{L}{t1}', f'=SUM({L}{r0}:{L}{last})', 'totm')
        put(ws, f'{L}{t2}', (f'=SUMIFS({CAH},{CAM},"支出",{CAN},{YM}+{L}$4)-SUMIFS({CAG},{CAM},"支出",{CAN},{YM}+{L}$4)') if m < 13 else f'=SUM(B{t2}:M{t2})', 'totm')
        put(ws, f'{L}{t3}', f'=ROUND({L}{t2}-{L}{t1},2)', 'totm')
    ws.freeze_panes = 'B5'


def s_pl(wb, d, rec):
    ws = wb[PL]
    head(ws, 4, [('A', '项目', 30), ('B', '说明', 34)] + [(CL(i), '', 12) for i in range(3, 16)], 'A076 简易利润表（按月，年份见首页）',
         '销售收入、销售成本来自管家婆销售出库（成本按库存表的加权单价）；平台费用、平台补贴来自拼多多店铺账务明细（按到账月份）；其他费用、其他收入来自资金台帐。'
         '下面「参考」两行：拼多多店铺回款 vs 管家婆里拼多多渠道的销售额，差得多说明拼多多卖出的货在管家婆没逐单出库（库存会偏多、成本偏少）。', last='O')
    month_hdr(ws, 4, 'C'); put(ws, 'O4', '全年合计', 'h')
    R = rec['R']
    rb = f"'{REC}'!$B${rec['rr'][0]}:$B${rec['r_oth']}"
    SAS, SAO, SAQ = RG(SA, 'S'), RG(SA, 'O'), RG(SA, 'Q')
    CAG, CAH, CAM, CAN = (RG(CA, c) for c in 'GHMN')

    def recsum(L, cls, sign=''):
        return f"={sign}SUMIF({rb},\"{cls}\",'{REC}'!{L}${rec['rr'][0]}:{L}${rec['r_oth']})"
    rows = [
        ('销售收入－拼多多', '管家婆销售出库，渠道＝拼多多', lambda L, ym: f'=SUMIFS({SAS},{SAO},"拼多多",{SAQ},{ym})'),
        ('销售收入－线下', '管家婆销售出库，渠道＝线下', lambda L, ym: f'=SUMIFS({SAS},{SAO},"线下",{SAQ},{ym})'),
        ('销售收入合计', '', lambda L, ym: f'={L}5+{L}6'),
        ('销售成本', '各月销售数量×加权单价（库存表）', None),
        ('毛利', '', lambda L, ym: f'={L}7-{L}8'),
        ('技术服务费净额', '基础＋百亿补贴－返还（账务明细）', lambda L, ym: recsum(L, '技术服务费', '-')),
        ('消费者体验提升计划', '账务明细', lambda L, ym: recsum(L, '体验计划', '-')),
        ('售后补偿（扣申诉补回）', '账务明细', lambda L, ym: recsum(L, '售后补偿', '-')),
        ('平台费用合计', '', lambda L, ym: f'={L}10+{L}11+{L}12'),
        ('平台补贴（补贴汇入）', '营销活动结算到账（账务明细）', lambda L, ym: recsum(L, '平台补贴')),
        ('其他费用', '资金台帐 属性＝支出', lambda L, ym: f'=SUMIFS({CAH},{CAM},"支出",{CAN},{ym})-SUMIFS({CAG},{CAM},"支出",{CAN},{ym})'),
        ('其他收入', '资金台帐 属性＝收入', lambda L, ym: f'=SUMIFS({CAG},{CAM},"收入",{CAN},{ym})-SUMIFS({CAH},{CAM},"收入",{CAN},{ym})'),
        ('净利润（简易）', '毛利－平台费用＋平台补贴－其他费用＋其他收入', lambda L, ym: f'={L}9-{L}13+{L}14-{L}15+{L}16'),
        (None, None, None),
        ('参考：拼多多店铺回款净额', '订单收入＋优惠券结算－退款（账务明细）', lambda L, ym: recsum(L, '店铺回款')),
        ('参考：差额', '店铺回款－管家婆拼多多销售额', lambda L, ym: f'={L}19-{L}5'),
    ]
    for i, (t, note, fn) in enumerate(rows):
        r = 5 + i
        if t is None:
            continue
        bold = t in ('销售收入合计', '毛利', '平台费用合计', '净利润（简易）')
        put(ws, f'A{r}', t, 'tot' if bold else 'lb'); put(ws, f'B{r}', note, 'lb')
        for m in range(1, 13):
            L = CL(2 + m)
            ym = f'{YM}+{L}$4'
            if t == '销售成本':
                k = CL(12 + m)
                f = f"=SUMPRODUCT('{INV}'!$H$6:$H$305,'{INV}'!{k}$6:{k}$305)"
            else:
                f = fn(L, ym)
            put(ws, f'{L}{r}', f, 'totm' if bold else 'am')
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})', 'totm' if bold else 'am')
    ws.freeze_panes = 'C5'


def s_home(wb, d, rec):
    ws = wb[HOME]
    for L, w in zip('ABCDEFG', (30, 16, 12, 30, 16, 40, 4)):
        ws.column_dimensions[L].width = w
    ws.merge_cells('A1:F1'); put(ws, 'A1', 'A076 拼多多电脑电商财务帐套 · 首页（报表本）', 't'); ws.row_dimensions[1].height = 28
    put(ws, 'A2', '录入本连接：', 'lb')
    ws.merge_cells('B2:F2')
    put(ws, 'B2', f"=IF(ISERROR(COUNTIF('{BS}'!$A$4:$A$5,\"?*\")),\"请先打开 A076-1_录入本.xlsx，再在这本按 F9／更新\",\"录入本已连上（数取自 A076-1_录入本.xlsx）\")", 'lb')
    red_if(ws, 'B2', 'LEFT($B$2,2)="请先"')
    green_if(ws, 'B2', 'LEFT($B$2,2)<>"请先"')
    ws.merge_cells('A3:F3')
    put(ws, 'A3', '用法：先开录入本（记账、粘管家婆和拼多多导出都在那本），再开这本看报表。两本放同一个文件夹，文件名别改。下面年份、月份可以改（淡黄格）。', 'tip')
    ws.row_dimensions[3].height = 32
    put(ws, 'A4', '年份', 'lb'); put(ws, 'B4', d['year'], 'i'); put(ws, 'C4', None, None)
    put(ws, 'D4', '月份（本月）', 'lb'); put(ws, 'E4', d['month'], 'i')
    dv_list(ws, '"2025,2026,2027,2028,2029,2030"', 'B4'); dv_list(ws, '"1,2,3,4,5,6,7,8,9,10,11,12"', 'E4')
    # 月份放 E4，D4 写说明：报表用 '首页'!$D$4 当月份 → 统一指 E4
    put(ws, 'A6', '资金账户', 'h'); put(ws, 'B6', '当前余额', 'h'); put(ws, 'D6', '本月（见上面月份）', 'h'); put(ws, 'E6', '金额', 'h')
    for i in range(6):
        r, b = 7 + i, ACC[0] + i
        put(ws, f'A{r}', f"=IF('{BS}'!$A${b}=\"\",\"\",'{BS}'!$A${b})", 'a')
        put(ws, f'B{r}', f"=IF(A{r}=\"\",\"\",'{BS}'!$D${b})", 'kpi')
    put(ws, 'A13', '合计', 'tot'); put(ws, 'B13', '=SUM(B7:B12)', 'totm')
    put(ws, 'C10', '含平台账务明细', 'lb')
    ym = "'首页'!$B$4*100+'首页'!$E$4"
    CAG, CAH, CAM, CAN = (RG(CA, c) for c in 'GHMN')
    pl = lambda r: f"=INDEX('{PL}'!$C${r}:$N${r},'首页'!$E$4)"
    kp = [('资金台帐收入（不含内部转账）', f'=SUMIFS({CAG},{CAN},{ym},{CAM},"<>内部转账")'),
          ('资金台帐支出（不含内部转账）', f'=SUMIFS({CAH},{CAN},{ym},{CAM},"<>内部转账")'),
          ('销售收入（管家婆）', pl(7)), ('销售成本', pl(8)), ('毛利', pl(9)), ('净利润（简易）', pl(17))]
    for i, (t, f) in enumerate(kp):
        put(ws, f'D{7 + i}', t, 'lb'); put(ws, f'E{7 + i}', f, 'kpi')
    put(ws, 'A15', '拼多多对账提示', 'h'); put(ws, 'B15', '笔数／单数', 'h'); put(ws, 'C15', '金额', 'h')
    ws.merge_cells('D15:F15'); put(ws, 'D15', '看哪里／怎么处理', 'h')
    R = rec['R']
    idx = lambda r: f"INDEX('{REC}'!$C${r}:$N${r},'首页'!$E$4)"
    A0, A1 = rec['A']
    B0, B1 = rec['B']
    items = [
        ('本月店铺净额 跟账务明细「#总计」差', '', f'={idx(R["diff"])}', '拼多多对账「一」；应为 0'),
        ('本月月末店铺余额（期初＋净额）', '', f'={idx(R["end"])}', '期初在录入本基础资料第 8 行'),
        ('本月提现没记进资金台帐的差额', '', f'={idx(R["cdiff"])}', '应为 0；提现在资金台帐记内部转账'),
        ('到账批次：对账单里没有明细的', f"=COUNTIF('{REC}'!$J${A0}:$J${A1},\"对账单里没有*\")",
         f"=SUMIFS('{REC}'!$E${A0}:$E${A1},'{REC}'!$J${A0}:$J${A1},\"对账单里没有*\")", '按批次号去后台下载这几批的对账单明细，粘进录入本'),
        ('对账单批次：账务明细还没到账的', f"=COUNTIF('{REC}'!$I${B0}:$I${B1},\"账务明细里还没到账*\")",
         f"=SUMIFS('{REC}'!$F${B0}:$F${B1},'{REC}'!$I${B0}:$I${B1},\"账务明细里还没到账*\")", '看结算时间，粘那个月的账务明细再看'),
        ('营销活动逐单不一致', f"='{REC}'!$I${rec['c0']}", '', '拼多多对账「三」'),
        ('百亿补贴技术服务费 扣了还没返', f"='{TF}'!$H$10", f"='{TF}'!$H$11", '技术服务费核对「三」；超过 30 天没返问客服'),
        ('返还账单里在账务明细找不到这单', f"='{TF}'!$C$10", f"='{TF}'!$C$11", '扣费在更早月份，把那个月账务明细也粘上'),
        ('返了又扣回（-t01）', f"='{TF}'!$C$7", f"='{TF}'!$C$8", '技术服务费核对「二」，问客服扣回原因'),
    ]
    for i, (t, n, a, how) in enumerate(items):
        r = 16 + i
        put(ws, f'A{r}', t, 'lb'); put(ws, f'B{r}', n if n else None, 'ai'); put(ws, f'C{r}', a if a else None, 'am')
        ws.merge_cells(f'D{r}:F{r}'); put(ws, f'D{r}', how, 'lbw')
    red_if(ws, 'B16:C24', 'AND(ISNUMBER(B16),B16<>0)')
    put(ws, 'A27', '跳转', 'sec')
    for i, sh in enumerate(REPORT[1:]):
        c = put(ws, f'A{28 + i}', f'→ {sh}', 'lnk')
        c.hyperlink = Hyperlink(ref=f'A{28 + i}', location=f"'{sh}'!A1")
    c = put(ws, 'D28', '→ 打开录入本 A076-1_录入本.xlsx（同一文件夹）', 'lnk')
    c.hyperlink = WB1
    ws.freeze_panes = 'A5'


HELP_TXT = [
    ('两本怎么用', ''),
    ('1', 'A076-1_录入本.xlsx：基础资料、资金台帐（手记）、采购入库、销售出库（粘管家婆导出）、拼多多三张粘贴区。平时只开这一本，快。'),
    ('2', 'A076-2_报表本.xlsx：首页、拼多多对账、技术服务费核对、库存、往来、费用统计、利润表。公式直接取录入本的数。'),
    ('3', '两本放同一个文件夹，文件名别改。看报表：先开录入本，再开报表本（问「更新链接」点更新），首页 B2 显示「录入本已连上」就对了。'),
    ('4', '录入本没开的时候，报表本别点「更新」：SUMIFS 取关着的文件会变成 #VALUE!。没开就先看上次存的数，要最新的先开录入本再按 F9。'),
    ('资金台帐', ''),
    ('5', '招商银行、微信、支付宝、拼多多店铺余额混在一张表记，一行一笔，按日期先后往下记；每行自动算这个账户记完这笔后的余额，「问题提示」列有字就是那行没填对。'),
    ('6', '拼多多提现到招行：账户＝拼多多店铺余额，收支类别＝内部转账，支出＝提现金额，对方账户＝招商银行。微信、支付宝转招行同理。'),
    ('7', '店铺里的订单收入、退款、技术服务费这些不用手记，粘账务明细就自动算进拼多多店铺余额。'),
    ('管家婆导出', ''),
    ('8', '采购：管家婆「商品采购明细查询」导出，从明细第 1 行复制到最后一行，粘到录入本「采购入库」A4。销售：「销售明细查询-按商品统计」同样粘到「销售出库」A4。'),
    ('9', '新商品、新客户、新供应商要先加到基础资料，不然库存和往来算不到（问题提示列会提醒）。'),
    ('10', '库存期初＝你给的管家婆库存表（你说不太准）。盘点后改基础资料的期初数量、期初单价。成本按「（期初金额＋采购金额）÷（期初数量＋采购数量）」加权平均。'),
    ('拼多多对账（核心）', ''),
    ('11', '每月从拼多多后台导出三张：店铺账务明细、营销活动结算对账单、技术服务费返还账单，分别粘到三张粘贴表 A4，下个月接着往下粘。'),
    ('12', '营销活动对账单对不上，是因为对账单按「结算时间」出、账务明细按「到账」出，而且账务明细里补贴汇入只有一笔批次总额。要按批次号对：账务明细备注里的批次号 → 对账单里同批次号的合计 → 应等于到账金额。'),
    ('13', '逐单：结算总金额＝成本金额（报名价）－商品售价（活动价）＋备注里的补贴优惠券。'),
    ('14', '技术服务费：百亿补贴单扣 0.6% 基础费＋1.4% 百亿费，1.4% 过一段时间按批次返还。「技术服务费核对」逐单列出扣了没返、返了又扣回（-t01）的单。'),
    ('15', '每月净额（收入合计＋支出合计）要等于账务明细最后一行「#总计」；期初店铺余额填在基础资料第 8 行。'),
    ('以后上管家婆财务', ''),
    ('16', '管家婆财务弄好之前先用这套；弄好后资金台帐可以停，拼多多对账这部分照样能用（粘导出就行）。'),
]


def s_help(wb, d):
    ws = wb[HELP]
    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 110
    ws.merge_cells('A1:B1'); put(ws, 'A1', 'A076 使用说明', 't'); ws.row_dimensions[1].height = 28
    r = 3
    for k, t in HELP_TXT:
        if not t:
            put(ws, f'A{r}', k, 'sec')
        else:
            put(ws, f'A{r}', k, 'txt'); put(ws, f'B{r}', t, 'txt')
            ws.row_dimensions[r].height = 30 if len(t) > 60 else 18
        r += 1


# ─────────────────────────── 一本版 ───────────────────────────
def build(d):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _styles(wb)
    colors = {BS: '595959', CA: '2F75B5', PU: '7030A0', SA: '7030A0', P1: 'C55A11', P2: 'C55A11', P3: 'C55A11'}
    for n in REPORT + ENTRY:
        ws = wb.create_sheet(n)
        ws.sheet_properties.tabColor = colors.get(n, '1F4E79' if n == HOME else '548235')
        ws.sheet_view.showGridLines = n not in (HOME, HELP)
    s_base(wb, d); s_cash(wb, d); s_pur(wb, d); s_sal(wb, d); s_p1(wb, d); s_p2(wb, d); s_p3(wb, d)
    rec = s_rec(wb, d)
    s_tf(wb, d); s_inv(wb, d); s_ar(wb, d); s_exp(wb, d); s_pl(wb, d, rec); s_home(wb, d, rec); s_help(wb, d)
    wb.active = 0
    wb.calculation.fullCalcOnLoad = False
    wb.calculation.calcId = 191029
    return wb, rec
