# -*- coding: utf-8 -*-
"""门店余额核对 · 生成脚本

在你的《多帐户收支登记表》里加两张表，原来的 27 张表一个字都不动：
  【门店余额核对】     门店余额 ＋ 收款 － 付款 ＋ 非门店年初数 ＝ 应存；应存跟 14 个账户的余额（实存）比，差额正常是 0
  【门店余额核对辅助】（隐藏）【数据录入】每一行算几个标记：计入日期、账户类型、是不是门店、是不是调拨、所属月份类型

【数据录入】可以照常插行、删行：辅助表按行号取数（INDEX(数据录入!$B:$B,ROW())），新表的求和范围也是按行号拼出来的
（INDEX(数据录入!$F:$F,5):INDEX(数据录入!$F:$F,30191)），插行删行都不会错位。
范围从第 5 行（账户期初）就开始取：期初那几行收入、支出是空的，算进来不影响数，上面删一行也不会把第一笔流水挤出去。

做法：openpyxl 只用来画新表（存成一个临时工作簿），再把新表的 XML 和它用到的样式直接塞进原文件 ——
原文件不经 openpyxl 存盘，【操作流程】里的 15 张截图、5 处批注、WPS 的动态数组都原样保留。
新表的数先在副本上用 LibreOffice 整本算一遍，再把算出来的数写进新表当缓存（打开就有数，改了数据录入照常自动重算）。

跑法：python3 build.py [原表] [输出]
"""
import copy
import glob
import os
import re
import shutil
import subprocess
import sys
import zipfile

from lxml import etree
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, '参考', '原表_数据录入_2026.9.27.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(ROOT, '多帐户收支登记表_加门店余额核对.xlsx')
TMP = os.path.join(HERE, '_tmp')

SH, AUX = '门店余额核对', '门店余额核对辅助'
D0, D1 = 5, 30191                   # 数据录入：第 5～18 行是各账户的期初（收入支出空着），第 19 行起是流水；取到 30191 行（跟各账户子表一样）
M0, M1 = 5, 1004                    # 门店月度收支统计：第 5 行起每个门店 12 行（留到 1004 行）
NS_DEFAULT = [('公司', '公司费用、利息、税费等'), ('配送', '进货付款、收调拨等（配送中心）'), ('预收', ''),
              ('收款', ''), ('其他收款', ''), ('股东借款', '股东借进借出'),
              ('内部转款', '自己账户之间转钱、支取现金：转出、转入两边的 I 列都要填「内部转款」'),
              ('房租及其他', ''), ('12月门店', '去年 12 月的门店款，没分到具体门店')]

NSX = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
NS = '{%s}' % NSX
RNS = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

# 【数据录入】的范围按行号拼：插行、删行都不会改变大小，跟辅助表永远一行对一行
DR = lambda c: f'INDEX(数据录入!${c}:${c},{D0}):INDEX(数据录入!${c}:${c},{D1})'
XR = lambda c: f'{AUX}!${c}${D0}:${c}${D1}'
MR = lambda c: f'门店月度收支统计!${c}${M0}:${c}${M1}'
CUT = '$L$4'                        # 截止日期（数字）：放在本表隐藏的 L4，复制出来的表也按自己的 B4 算
LAST, BANK0, BOOKY = (f'{AUX}!$M${i}' for i in range(2, 5))
MY = '门店月度收支统计!$B$2'          # 那张表统计的年份（首月手加数按它算）

# ---- 版面（行号） ----
R_CUT = 4
R1 = dict(store=9, recv=10, pay=11, co=12, co_a=13, co_b=14, co_c=15, should=16, bank=17, diff=18, stmt=19)
M_H = 21; M_0 = M_H + 2                                   # ② 按月：23～34
X_H = 36; X_0 = X_H + 2                                   # ③ 差额从哪来：38～
X = dict(a=X_0, b=X_0 + 1, c=X_0 + 2, d=X_0 + 3, sub=X_0 + 4, eff=X_0 + 5, h3=X_0 + 6, inner=X_0 + 7,
         nodate=X_0 + 8, chk=X_0 + 9)
N_H = 49; N_0 = N_H + 2; N_1 = N_0 + 14                   # ④ 不算门店的名字 51～65
N_PREV, N_BL, N_TOT = N_1 + 1, N_1 + 2, N_1 + 3           # 66 门店上年的款，67 没填门店，68 合计
S_H = 70; S_0 = S_H + 2; S_1 = S_0 + 79                   # ⑤ 门店 72～151（80 家）
S_RES, S_TOT = S_1 + 1, S_1 + 2
A_H = S_TOT + 2; A_0 = A_H + 2; A_1 = A_0 + 19            # ⑥ 账户 20 行
A_TOT = A_1 + 1
NSR = f'$B${N_0}:$B${N_1}'
LAYOUT = dict(R1=R1, M_0=M_0, X=X, N_0=N_0, N_1=N_1, N_PREV=N_PREV, N_BL=N_BL, N_TOT=N_TOT, S_0=S_0, S_1=S_1,
              S_RES=S_RES, S_TOT=S_TOT, A_0=A_0, A_1=A_1, A_TOT=A_TOT)

YH = '微软雅黑'
F_TITLE = Font(name=YH, size=15, bold=True, color='FF1F3864')
F_SEC = Font(name=YH, size=11, bold=True, color='FFFFFFFF')
F_HDR = Font(name=YH, size=10, bold=True, color='FF1F3864')
F_TXT = Font(name=YH, size=10)
F_TXTB = Font(name=YH, size=10, bold=True)
F_SUB = Font(name=YH, size=9, color='FF404040')
F_NOTE = Font(name=YH, size=9, color='FF595959')
F_BIG = Font(name=YH, size=12, bold=True, color='FF1F3864')
F_IN = Font(name=YH, size=10, bold=True, color='FF0000C0')
FILL_SEC = PatternFill('solid', fgColor='FF2F5597')
FILL_HDR = PatternFill('solid', fgColor='FFD9E1F2')
FILL_IN = PatternFill('solid', fgColor='FFFFF2CC')
FILL_AUTO = PatternFill('solid', fgColor='FFF2F2F2')
FILL_KEY = PatternFill('solid', fgColor='FFE2EFDA')
FILL_TOT = PatternFill('solid', fgColor='FFDDEBF7')
CF = lambda c: PatternFill(fill_type='solid', fgColor=c, bgColor=c)      # 条件格式的底色：前景背景都写上，Excel / WPS 都认
THIN = Side(style='thin', color='FFBFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
AL = Alignment(horizontal='left', vertical='center', wrap_text=True)
ALI = Alignment(horizontal='left', vertical='center', wrap_text=True, indent=2)
AC = Alignment(horizontal='center', vertical='center', wrap_text=True)
AR = Alignment(horizontal='right', vertical='center')
MONEY = '#,##0.00;[Red]-#,##0.00'
DATE = 'yyyy-mm-dd'
UNLOCK = Protection(locked=False)


def put(ws, ref, v=None, font=F_TXT, fill=None, fmt=None, align=AL, border=BOX):
    c = ws[ref]
    c.value = v
    c.font = font
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    c.alignment = align
    if border is not None:
        c.border = border
    return c


def merge(ws, rng, v=None, **kw):
    from openpyxl.utils import column_index_from_string as ci, get_column_letter as cl
    first = rng.split(':')[0]
    put(ws, first, v, **kw)
    a, b = rng.split(':')
    col0, col1 = re.match(r'[A-Z]+', a).group(), re.match(r'[A-Z]+', b).group()
    row = int(re.search(r'\d+', a).group())
    for k in range(ci(col0) + 1, ci(col1) + 1):
        put(ws, f'{cl(k)}{row}', None, font=kw.get('font', F_TXT), fill=kw.get('fill'), border=kw.get('border', BOX))
    ws.merge_cells(rng)


def section(ws, row, text, last='J'):
    merge(ws, f'A{row}:{last}{row}', text, font=F_SEC, fill=FILL_SEC, border=None)
    ws.row_dimensions[row].height = 22


def heads(ws, row, items, height=30):
    for col, t in items:
        put(ws, f'{col}{row}', t, font=F_HDR, fill=FILL_HDR, align=AC)
    ws.row_dimensions[row].height = height


def le(r):
    return f'"<="&{r}'


def cut_sum(fld, *crit, cut=CUT):
    """数据录入某一列（F 收入 / G 支出），在截止日以前、满足条件的合计"""
    return f'SUMIFS({DR(fld)},{",".join(crit + (XR("A"), le(cut)))})'


# ============================================================ 画【门店余额核对】
def build_main():
    wb = Workbook()
    ws = wb.active
    ws.title = SH
    ws.sheet_properties.tabColor = 'FF00B050'
    ws.sheet_view.showGridLines = False
    for col, w in dict(A=9, B=42, C=17, D=17, E=17, F=17, G=17, H=19, I=17, J=14, K=4, L=4).items():
        ws.column_dimensions[col].width = w
    r = R1
    merge(ws, 'A1:J1', '门店余额核对　　门店余额 ＋ 收款 － 付款 ＋ 非门店年初数 ＝ 银行余额', font=F_TITLE, border=None)
    ws.row_dimensions[1].height = 30
    merge(ws, 'A2:J2',
          '全部从【数据录入】自动取数，只有黄格子可以填。门店余额＝每个门店的年初余额（取自【门店月度收支统计】）＋首月手加＋本年收入－支出；'
          '收款、付款＝不算门店的收支（见 ④）；非门店年初数＝年初各账户余额－门店年初余额（含首月手加），是负数说明门店账上的数比银行里的钱多（见 4a～4c）。'
          '领导的公式从开账算起，这本表从账本年初算起，所以单列一行。调拨、货款直抵费用、平台款分摊这些没走银行账户的记录收付应该相抵，不相抵的就是差额（见 ③）。'
          '原来的表一格都没动；【数据录入】插行、删行都没关系。',
          font=F_NOTE, border=None)
    ws.row_dimensions[2].height = 50
    put(ws, f'A{R_CUT}', '截止日期', font=F_TXTB, fill=FILL_HDR, align=AC)
    c = put(ws, f'B{R_CUT}', None, font=F_IN, fill=FILL_IN, fmt=DATE, align=AC)
    c.protection = UNLOCK
    merge(ws, f'C{R_CUT}:F{R_CUT}', '← 空着＝全部（到最后一笔）；填某一天＝算到那天为止（没写日期的记录按「收支所属月份」的月底算）',
          font=F_NOTE, border=None)
    put(ws, f'G{R_CUT}', '现在算的是', font=F_TXTB, fill=FILL_HDR, align=AC)
    merge(ws, f'H{R_CUT}:J{R_CUT}', f'=IF(ISNUMBER($B${R_CUT}),"截至 "&TEXT($B${R_CUT},"yyyy-mm-dd"),'
                                    f'"全部（最后一笔 "&TEXT({LAST},"yyyy-mm-dd")&"）")', font=F_TXTB, fill=FILL_KEY)
    ws.row_dimensions[R_CUT].height = 24
    put(ws, 'A5', '账本年初', font=F_TXTB, fill=FILL_HDR, align=AC)
    put(ws, 'B5', '=数据录入!$B$5', font=F_TXT, fill=FILL_AUTO, fmt=DATE, align=AC)
    acc_n = 'ROWS(数据录入!$C$5:$C$18)'
    merge(ws, 'C5:J5', f'="这一天 "&{acc_n}&" 个账户的余额合计 "&TEXT({BANK0},"#,##0.00")&" 元（【数据录入】第 5～"&(4+{acc_n})&" 行的期初）"',
          font=F_NOTE, border=None)
    put(ws, f'L{R_CUT}', f'=IF(ISNUMBER($B${R_CUT}),$B${R_CUT},2958465)', font=F_NOTE, border=None)    # 截止日期（数字）
    dv = DataValidation(type='date', operator='greaterThan', formula1='36526', allow_blank=True,
                        showErrorMessage=True, errorTitle='截止日期', error='填日期，比如 2026/9/30；空着＝全部')
    ws.add_data_validation(dv)
    dv.add(f'B{R_CUT}')

    # ---------------- ① 核对 ----------------
    section(ws, 7, '① 核对：门店余额 ＋ 收款 － 付款 ＋ 非门店年初数 ＝ 应存，应存 ＝ 各账户的账上余额（实存）')
    heads(ws, 8, [('A', '序'), ('B', '项目'), ('C', '金额'), ('D', '说明')], 22)
    merge(ws, 'D8:J8', '说明', font=F_HDR, fill=FILL_HDR, align=AC)
    store_flow = f'{cut_sum("F", XR("C"), "1")}-{cut_sum("G", XR("C"), "1")}'
    rows1 = [
        (r['store'], '1', '门店余额（各门店：年初 ＋ 首月手加 ＋ 收入 － 支出）', f'=$C${S_TOT}+$D${S_TOT}+{store_flow}',
         '每个门店多少看 ⑤。门店＝【门店月度收支统计】里的门店（含自助药机、疼痛馆、关了的店），不含 ④ 里那些名字。'
         '注意：银行到账了、I 列还没填门店的，先算在收款里，门店余额会偏小（差额照样是 0，看不出来）'),
        (r['recv'], '2', '加：收款（不算门店的收入）', f'={cut_sum("F", XR("C"), chr(34) + "<>1" + chr(34))}',
         '公司、配送、内部转款、股东借款、没填门店的、门店上年的款……分别多少看 ④（含调拨等不走银行的，所以比银行流水大）'),
        (r['pay'], '3', '减：付款（不算门店的支出）', f'={cut_sum("G", XR("C"), chr(34) + "<>1" + chr(34))}', '同上'),
        (r['co'], '4', '加：非门店年初数（＝ 4a ＋ 4b ＋ 4c）', f'=ROUND($C${r["co_a"]}+$C${r["co_b"]}+$C${r["co_c"]},2)',
         '＝年初各账户余额 － 年初门店余额（含首月手加的数）。是倒推出来的数，下面三行是它的来历'),
        (r['co_a'], '4a', '其中：配送、公司等的年初数（取自【门店月度收支统计】）', f'=ROUND(SUM($L${N_0}:$L${N_1}),2)',
         '④ 里那些名字在【门店月度收支统计】G 列的年初数（现在是配送、公司）'),
        (r['co_b'], '4b', '其中：年初没对平的（倒挤，要查明）', f'=ROUND({BANK0}-$C${S_TOT}-$C${r["co_a"]},2)',
         '＝年初银行余额 －（门店年初 ＋ 配送公司年初）。旧表年初就差这么多，比如年底在路上还没存进银行的门店款；要跟领导说明'),
        (r['co_c'], '4c', '其中：减去门店首月手加的数', f'=-ROUND($D${S_TOT},2)',
         '【门店月度收支统计】首月公式后面手加的数（⑤ D 列），【数据录入】里没有这笔钱，加在门店上就要在这里减掉'),
        (r['should'], '＝', '应存（1 ＋ 2 － 3 ＋ 4）', f'=ROUND($C${r["store"]}+$C${r["recv"]}-$C${r["pay"]}+$C${r["co"]},2)',
         f'=IF(ISNUMBER($B${R_CUT}),"","＝【数据录入】最上面 H3 的「账户余额」"&TEXT(数据录入!$H$3,"#,##0.00")&'
         f'"（没筛选时）。那个数把不走银行的记录也算进去了，所以不等于银行余额；银行余额看下一行")'),
        (r['bank'], '5', '实存（各账户的账上余额合计）',
         f'=ROUND({BANK0}+{cut_sum("F", XR("B"), "1")}-{cut_sum("G", XR("B"), "1")},2)',
         '账户＝【数据录入】第 5～18 行那些；每个账户多少看 ⑥（跟各账户子表的「当前余额」一个算法）'),
        (r['diff'], '差', '差额（实存 － 应存，正常是 0）', f'=ROUND($C${r["bank"]}-$C${r["should"]},2)', None),
        (r['stmt'], '6', '银行对账单余额合计（⑥ G 列填了的按对账单，没填的按账上）',
         f'=IF(SUM($J${A_0}:$J${A_1})=0,"",ROUND(SUM($I${A_0}:$I${A_1}),2))', None),
    ]
    for rr, no, lab, f, note in rows1:
        key = rr in (r['should'], r['bank'], r['diff'])
        sub = rr in (r['co_a'], r['co_b'], r['co_c'])
        put(ws, f'A{rr}', no, font=F_SUB if sub else F_TXTB, fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', lab, font=F_SUB if sub else (F_TXTB if key else F_TXT), fill=FILL_KEY if key else None,
            align=ALI if sub else AL)
        put(ws, f'C{rr}', f, font=F_BIG if key else (F_SUB if sub else F_TXTB), fill=FILL_KEY if key else FILL_AUTO,
            fmt=MONEY, align=AR)
        if note is not None:
            merge(ws, f'D{rr}:J{rr}', note, font=F_NOTE)
        ws.row_dimensions[rr].height = 44 if rr == r['store'] else (
            30 if (sub or rr in (r['co'], r['should'], r['diff'], r['stmt'], r['recv'])) else 22)
    merge(ws, f'D{r["diff"]}:J{r["diff"]}',
          f'=IF(ABS($C${r["diff"]})<0.005,"✓ 账面对上了：门店余额＋收款－付款＋非门店年初数＝各账户的账上余额（跟银行实际余额核对：在 ⑥ 填对账单余额）",'
          f'"✗ 应存比实存"&IF($C${r["diff"]}<0,"多 ","少 ")&TEXT(ABS($C${r["diff"]}),"#,##0.00")&" 元：多半是不走银行的记录记了两遍、或漏了冲总数（看 ③），'
          f'不是银行"&IF($C${r["diff"]}<0,"少","多")&"了钱")',
          font=F_TXTB)
    merge(ws, f'D{r["stmt"]}:J{r["stmt"]}',
          f'=IF($C${r["stmt"]}="","⑥ G 列还没填银行对账单余额（填了就能跟银行的实际余额对）",'
          f'IF(ABS($C${r["stmt"]}-$C${r["bank"]})<0.005,"✓ 对账单跟账上余额一致",'
          f'"✗ 对账单比账上"&IF($C${r["stmt"]}<$C${r["bank"]},"少 ","多 ")&TEXT(ABS($C${r["stmt"]}-$C${r["bank"]}),"#,##0.00")&" 元：看 ⑥ 是哪个账户"))',
          font=F_TXTB)
    ws.conditional_formatting.add(f'C{r["diff"]}:J{r["diff"]}', FormulaRule(
        formula=[f'ABS($C${r["diff"]})>=0.005'], fill=CF('FFFFC7CE'), font=Font(color='FF9C0006', bold=True)))
    ws.conditional_formatting.add(f'D{r["diff"]}:J{r["diff"]}', FormulaRule(
        formula=[f'ABS($C${r["diff"]})<0.005'], font=Font(color='FF00803C', bold=True)))

    # ---------------- ② 按月 ----------------
    section(ws, M_H, '② 按月看（按日期算到每个月月底；【门店月度收支统计】按所属月份算，两边同一个月的门店余额对不上是正常的）')
    heads(ws, M_H + 1, [('A', '月份'), ('B', '月底'), ('C', '门店余额'), ('D', '收款（年初起累计）'),
                        ('E', '付款（年初起累计）'), ('F', '非门店年初数'), ('G', '应存'), ('H', '实存（账上）'),
                        ('I', '差额'), ('J', '')])
    for k in range(1, 13):
        rr = M_0 + k - 1
        show = f'DATE({BOOKY},{k},1)<={LAST}'
        b = f'$B{rr}'
        put(ws, f'A{rr}', f'{k}月', font=F_TXTB, fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', f'=EOMONTH(DATE({BOOKY},{k},1),0)', fill=FILL_AUTO, fmt=DATE, align=AC)
        fs = {
            'C': f'$C${S_TOT}+$D${S_TOT}+{cut_sum("F", XR("C"), "1", cut=b)}-{cut_sum("G", XR("C"), "1", cut=b)}',
            'D': cut_sum('F', XR('C'), '"<>1"', cut=b),
            'E': cut_sum('G', XR('C'), '"<>1"', cut=b),
            'F': f'$C${r["co"]}',
            'G': f'C{rr}+D{rr}-E{rr}+F{rr}',
            'H': f'{BANK0}+{cut_sum("F", XR("B"), "1", cut=b)}-{cut_sum("G", XR("B"), "1", cut=b)}',
            'I': f'H{rr}-G{rr}',
        }
        for col, f in fs.items():
            put(ws, f'{col}{rr}', f'=IF(NOT({show}),"",ROUND({f},2))', fill=FILL_AUTO, fmt=MONEY, align=AR,
                font=F_TXTB if col == 'I' else F_TXT)
        put(ws, f'J{rr}', f'=IF(I{rr}="","",IF(ABS(I{rr})<0.005,"✓","✗"))', font=F_TXTB, align=AC)
        ws.row_dimensions[rr].height = 18
    ws.conditional_formatting.add(f'I{M_0}:J{M_0 + 11}', FormulaRule(
        formula=[f'AND($I{M_0}<>"",ABS(N($I{M_0}))>=0.005)'], fill=CF('FFFFC7CE'), font=Font(color='FF9C0006', bold=True)))
    ws.conditional_formatting.add(f'J{M_0}:J{M_0 + 11}', FormulaRule(
        formula=[f'$J{M_0}="✓"'], font=Font(color='FF00803C', bold=True)))

    # ---------------- ③ 差额从哪来 ----------------
    section(ws, X_H, '③ 差额从哪来：没走银行账户的记录，收入和支出应该正好相抵')
    heads(ws, X_H + 1, [('A', '序'), ('B', '项目'), ('C', '笔数'), ('D', '收入'), ('E', '支出'), ('F', '净额（应为 0）')])
    merge(ws, f'G{X_H + 1}:J{X_H + 1}', '说明', font=F_HDR, fill=FILL_HDR, align=AC)
    cnt = lambda *crit: f'SUMIFS({XR("E")},{",".join(crit + (XR("A"), le(CUT)))})'
    zd = (DR('C'), '"货款直抵费用"')
    xrows = [
        ('a', '没记账户的调拨（账户空着、摘要里有「调拨」）',
         cnt(XR('B'), '2', XR('D'), '1'), cut_sum('F', XR('B'), '2', XR('D'), '1'), cut_sum('G', XR('B'), '2', XR('D'), '1'),
         '配送收调拨、门店付调拨，一收一付应该相等'),
        ('b', '没记账户的其他记录（平台款、保险理赔款分到门店，报销拆分等）',
         f'{cnt(XR("B"), "2")}-C{X["a"]}', f'{cut_sum("F", XR("B"), "2")}-D{X["a"]}', f'{cut_sum("G", XR("B"), "2")}-E{X["a"]}',
         '常见原因：银行那一行已经记了（I 列空着＝算在收款、付款里），又用账户空着的行按门店分了一遍，同一笔钱记了两次。'
         '改法二选一：① 补一行账户空、I 列空的「冲总数」，方向跟分摊行相反（分摊记收入就补支出，分摊记支出就补收入），金额＝分摊合计；'
         '② 删掉分摊行，直接把银行那一行拆开填门店'),
        ('c', '货款直抵费用（门店用营业款直接付的费用）',
         cnt(*zd), cut_sum('F', *zd), cut_sum('G', *zd), '门店拿营业款付了费用：收入记营业款、支出记费用，两边应该相等'),
        ('d', '账户不是【数据录入】第 5～18 行那些银行现金账户的（虚拟账户、预收款项、聚合等，或写错的）',
         f'{cnt(XR("B"), "0")}-C{X["c"]}', f'{cut_sum("F", XR("B"), "0")}-D{X["c"]}', f'{cut_sum("G", XR("B"), "0")}-E{X["c"]}',
         '只有【数据录入】第 5～18 行那些账户算银行；别的账户名（基础资料 A 列下面那些，或写错的）不走银行，收付也要配平'),
    ]
    for no, lab, fc, fi, fe, note in xrows:
        rr = X[no]
        put(ws, f'A{rr}', no, font=F_TXTB, fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', lab)
        put(ws, f'C{rr}', f'={fc}', fill=FILL_AUTO, fmt='0', align=AR)
        put(ws, f'D{rr}', f'={fi}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{rr}', f'={fe}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{rr}', f'=ROUND(D{rr}-E{rr},2)', font=F_TXTB, fill=FILL_AUTO, fmt=MONEY, align=AR)
        merge(ws, f'G{rr}:J{rr}', note, font=F_NOTE)
        ws.row_dimensions[rr].height = 72 if no == 'b' else 32
    rs = X['sub']
    put(ws, f'A{rs}', '', fill=FILL_TOT)
    put(ws, f'B{rs}', '小计：没走银行账户的记录（净额应为 0）', font=F_TXTB, fill=FILL_TOT)
    for col in 'CDEF':
        put(ws, f'{col}{rs}', f'=SUM({col}{X["a"]}:{col}{X["d"]})', font=F_TXTB, fill=FILL_TOT,
            fmt='0' if col == 'C' else MONEY, align=AR)
    merge(ws, f'G{rs}:J{rs}', '这一行不是 0，差额就不是 0：去【数据录入】按「公司账户」筛选空白（或货款直抵费用），按日期、摘要找没配平的那几笔',
          font=F_NOTE, fill=FILL_TOT)
    ws.row_dimensions[rs].height = 30
    re_ = X['eff']
    put(ws, f'A{re_}', '', fill=FILL_TOT)
    put(ws, f'B{re_}', '对 ① 差额的影响（＝小计取反：不走银行的记录多记了多少，应存就虚高多少；不是银行少了钱）', font=F_TXTB, fill=FILL_TOT)
    for col in 'CDE':
        put(ws, f'{col}{re_}', None, fill=FILL_TOT)
    put(ws, f'F{re_}', f'=ROUND(-F{rs},2)', font=F_TXTB, fill=FILL_TOT, fmt=MONEY, align=AR)
    merge(ws, f'G{re_}:J{re_}', f'=IF(ABS(F{re_}-$C${r["diff"]})<0.005,"✓ 跟 ① 的差额一致","✗ 跟 ① 的差额不一致（公式被改动了？）")',
          font=F_TXTB, fill=FILL_TOT)
    ws.row_dimensions[re_].height = 44
    ws.conditional_formatting.add(f'F{X["a"]}:F{rs}', FormulaRule(formula=[f'ABS(N(F{X["a"]}))>=0.005'], fill=CF('FFFFC7CE'),
                                                                  font=Font(color='FF9C0006', bold=True)))
    for key, lab, cval, fval, note in [
        ('h3', '【数据录入】H3 的「账户余额」', None, '=IF(ISNUMBER($B$4),"",数据录入!$H$3)',
         f'=IF(ISNUMBER($B$4),"（填了截止日期就不比）",IF(ABS(F{X["h3"]}-$C${r["should"]})<0.005,"＝ ① 的应存（不是银行余额）",'
         f'"跟应存不一样：数据录入可能开着筛选（H3 只算筛选出来的行）"))'),
        ('inner', '内部转款净额（应为 0）', None,
         f'=ROUND({cut_sum("F", DR("I"), chr(34) + "内部转款" + chr(34))}-{cut_sum("G", DR("I"), chr(34) + "内部转款" + chr(34))},2)',
         '转出、转入两边的 I 列都要填「内部转款」；不是 0 多半是转入那笔 I 列空着、或者支取现金只记了一边。不影响差额，但 ④ 里公司、没填门店的数会不准'),
        ('nodate', '日期空着或不是日期、所属月份也没填或填错的记录（笔数）', f'=SUMIFS({XR("E")},{XR("A")},0)', None,
         '这些按年初算进去了（截止日期填哪天都算上）。去【数据录入】把日期改成 2026/9/15 这样的日期，或者填上所属月份'),
    ]:
        rr = X[key]
        put(ws, f'A{rr}', '参考', font=F_TXTB, fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', lab)
        put(ws, f'C{rr}', cval, fill=FILL_AUTO, fmt='0', align=AR)
        for col in 'DE':
            put(ws, f'{col}{rr}', None, fill=FILL_AUTO)
        put(ws, f'F{rr}', fval, fill=FILL_AUTO, fmt=MONEY, align=AR)
        merge(ws, f'G{rr}:J{rr}', note, font=F_NOTE)
        ws.row_dimensions[rr].height = 32
    ws.conditional_formatting.add(f'F{X["inner"]}', FormulaRule(formula=[f'ABS(N(F{X["inner"]}))>=0.005'], fill=CF('FFFCE4D6')))
    rk = X['chk']
    put(ws, f'A{rk}', '自查', font=F_TXTB, fill=FILL_HDR, align=AC)
    put(ws, f'B{rk}', '④⑤⑥ 明细合计跟 ① 一致吗；④ 名字有没有重复；⑤ 放不放得下')
    for col in 'CDEF':
        put(ws, f'{col}{rk}', None, fill=FILL_AUTO)
    merge(ws, f'G{rk}:J{rk}',
          f'=IF(MAX($K${N_0}:$K${N_1})>1,"✗ ④ 里有重复的名字，删掉一个",'
          f'IF(MAX({AUX}!$J${M0}:$J${M1})>{S_1 - S_0 + 1},"✗ 门店超过 {S_1 - S_0 + 1} 家，⑤ 放不下，要请人加行",'
          f'IF(AND(ABS($G${S_TOT}-$C${r["store"]})<0.005,ABS($C${N_TOT}-$C${r["recv"]})<0.005,'
          f'ABS($D${N_TOT}-$C${r["pay"]})<0.005,ABS($F${A_TOT}-$C${r["bank"]})<0.005),"✓ 一致",'
          f'"✗ 不一致（公式被改动了？）")))', font=F_TXTB)

    # ---------------- ④ 收款 / 付款 明细 ----------------
    section(ws, N_H, '④ 收款 / 付款 明细（这些名字不算门店；B 列黄格可以改、可以往下加）')
    heads(ws, N_H + 1, [('A', '序'), ('B', '名字（数据录入 I 列「公司」）'), ('C', '收款'), ('D', '付款'), ('E', '净额'),
                        ('F', '其中走银行的·收'), ('G', '其中走银行的·付'), ('H', '说明')])
    merge(ws, f'H{N_H + 1}:J{N_H + 1}', '说明', font=F_HDR, fill=FILL_HDR, align=AC)
    for i in range(N_1 - N_0 + 1):
        rr = N_0 + i
        nm, note = NS_DEFAULT[i] if i < len(NS_DEFAULT) else (None, '')
        b = f'$B{rr}'
        nmc = (DR('I'), b)
        put(ws, f'A{rr}', i + 1, fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', nm, font=F_IN, fill=FILL_IN).protection = UNLOCK
        put(ws, f'C{rr}', f'=IF({b}="","",{cut_sum("F", *nmc)})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'D{rr}', f'=IF({b}="","",{cut_sum("G", *nmc)})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{rr}', f'=IF({b}="","",C{rr}-D{rr})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{rr}', f'=IF({b}="","",{cut_sum("F", *nmc, XR("B"), "1")})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{rr}', f'=IF({b}="","",{cut_sum("G", *nmc, XR("B"), "1")})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        merge(ws, f'H{rr}:J{rr}', note, font=F_NOTE)
        put(ws, f'K{rr}', f'=IF({b}="",0,COUNTIF({NSR},{b}))', font=F_NOTE, border=None)          # 重复几次
        put(ws, f'L{rr}', f'=IF({b}="",0,SUMIFS({MR("G")},{MR("B")},{b}))', font=F_NOTE, border=None)   # 门店月度收支统计里的年初数
        ws.row_dimensions[rr].height = 30 if note and len(note) > 22 else 18
    dvn = DataValidation(type='custom', formula1=f'COUNTIF({NSR},B{N_0})<=1', allow_blank=True, showErrorMessage=True,
                         errorTitle='名字重复了', error='这个名字上面已经有了')
    ws.add_data_validation(dvn)
    dvn.add(f'B{N_0}:B{N_1}')
    for rr, lab, crit, note in [
        (N_PREV, '门店上年的款（所属月份是去年的门店记录）', (XR('C'), '2'),
         '比如 1 月初存进来的去年 12 月营业款（所属月份填 202512）：年初余额里已经算过，不再加到门店上'),
    ]:
        put(ws, f'A{rr}', '', fill=FILL_HDR)
        put(ws, f'B{rr}', lab)
        put(ws, f'C{rr}', f'={cut_sum("F", *crit)}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'D{rr}', f'={cut_sum("G", *crit)}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{rr}', f'=C{rr}-D{rr}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{rr}', f'={cut_sum("F", *crit, XR("B"), "1")}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{rr}', f'={cut_sum("G", *crit, XR("B"), "1")}', fill=FILL_AUTO, fmt=MONEY, align=AR)
        merge(ws, f'H{rr}:J{rr}', note, font=F_NOTE)
        ws.row_dimensions[rr].height = 30
    rr = N_BL
    put(ws, f'A{rr}', '', fill=FILL_HDR)
    put(ws, f'B{rr}', '没填门店的（I 列空着）')
    for col, fld, extra in (('C', 'F', ()), ('D', 'G', ()), ('F', 'F', (XR('B'), '1')), ('G', 'G', (XR('B'), '1'))):
        put(ws, f'{col}{rr}', f'={cut_sum(fld, XR("C"), chr(34) + "<>1" + chr(34), *extra)}-SUM({col}{N_0}:{col}{N_PREV})',
            fill=FILL_AUTO, fmt=MONEY, align=AR)
    put(ws, f'E{rr}', f'=C{rr}-D{rr}', fill=FILL_AUTO, fmt=MONEY, align=AR)
    merge(ws, f'H{rr}:J{rr}', '银行到账没分门店的（如「24小时微信回款」）。注意：已经用没账户的行分到门店了，就不要再给银行那一行填门店，否则门店算两遍',
          font=F_NOTE)
    ws.row_dimensions[rr].height = 44
    rr = N_TOT
    put(ws, f'A{rr}', '', fill=FILL_TOT)
    put(ws, f'B{rr}', '合计（＝ ① 的收款 / 付款）', font=F_TXTB, fill=FILL_TOT)
    for col in 'CDEFG':
        put(ws, f'{col}{rr}', f'=SUM({col}{N_0}:{col}{N_BL})', font=F_TXTB, fill=FILL_TOT, fmt=MONEY, align=AR)
    merge(ws, f'H{rr}:J{rr}', '收款付款里含不走银行的（配送收调拨等），所以比银行流水的数大', font=F_NOTE, fill=FILL_TOT)

    # ---------------- ⑤ 门店余额明细 ----------------
    section(ws, S_H, '⑤ 门店余额明细（门店名单、年初余额取自【门店月度收支统计】）')
    heads(ws, S_H + 1, [('A', '序'), ('B', '门店'), ('C', '年初余额'), ('D', '首月手加的数'), ('E', '收入'),
                        ('F', '支出'), ('G', '期末余额'),
                        ('H', f'=IF({MY}={BOOKY},"所属月份没填的"&CHAR(10)&"（门店月度收支统计没算到）",'
                              f'"门店月度收支统计 B2 不是"&CHAR(10)&{BOOKY}&" 年，这列不算")'), ('I', '提示')], 40)
    merge(ws, f'I{S_H + 1}:J{S_H + 1}', '提示', font=F_HDR, fill=FILL_HDR, align=AC)
    for i in range(S_1 - S_0 + 1):
        rr = S_0 + i
        b = f'$B{rr}'
        put(ws, f'A{rr}', f'=IF({b}="","",{i + 1})', fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', f'=IFERROR(INDEX({AUX}!$H${M0}:$H${M1},MATCH({i + 1},{AUX}!$J${M0}:$J${M1},0)),"")')
        put(ws, f'C{rr}', f'=IF({b}="","",SUMIFS({MR("G")},{MR("B")},{b}))', fill=FILL_AUTO, fmt=MONEY, align=AR)
        # K：这个门店所属月份在那张表年份内的净额（＝那张表用公式取到的部分）
        put(ws, f'K{rr}', f'=IF({b}="","",SUMIFS({DR("F")},{DR("I")},{b},{XR("F")},0)-SUMIFS({DR("G")},{DR("I")},{b},{XR("F")},0))',
            font=F_NOTE, border=None)
        # 已关的店在那张表里只有单列的一行（没有 D、E 公式），首月手加按 0
        put(ws, f'D{rr}', f'=IF({b}="","",IF(COUNTIF({MR("B")},{b})<12,0,'
                          f'ROUND(SUMIFS({MR("D")},{MR("B")},{b})-SUMIFS({MR("E")},{MR("B")},{b})-K{rr},2)))',
            fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{rr}', f'=IF({b}="","",{cut_sum("F", DR("I"), b, XR("C"), "1")})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{rr}', f'=IF({b}="","",{cut_sum("G", DR("I"), b, XR("C"), "1")})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{rr}', f'=IF({b}="","",C{rr}+D{rr}+E{rr}-F{rr})', font=F_TXTB, fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'H{rr}', f'=IF(OR({b}="",{MY}<>{BOOKY}),"",ROUND(SUMIFS({DR("F")},{DR("I")},{b},{XR("C")},1,{XR("F")},"<>0")'
                          f'-SUMIFS({DR("G")},{DR("I")},{b},{XR("C")},1,{XR("F")},"<>0"),2))', fill=FILL_AUTO, fmt=MONEY, align=AR)
        merge(ws, f'I{rr}:J{rr}',
              f'=IF({b}="","",IF(AND(ABS(N(D{rr}))>=0.005,COUNTIF($D${S_0}:$D${S_1},D{rr})>1),'
              f'"首月手加的数跟别家一样，查一下",'
              f'IF(COUNTIF({MR("B")},{b})<12,IF(OR(ABS(N(E{rr}))>=0.005,ABS(N(F{rr}))>=0.005),"已关的店本年还有收支，核对",'
              f'"已关的店（统计表里单列一行）"),'
              f'IF(ABS(N(H{rr}))>=0.005,"所属月份没填，统计表没算到",""))))', font=F_NOTE)
    ws.conditional_formatting.add(f'H{S_0}:H{S_1}', FormulaRule(formula=[f'ABS(N(H{S_0}))>=0.005'], fill=CF('FFFCE4D6')))
    ws.conditional_formatting.add(f'D{S_0}:D{S_1}', FormulaRule(formula=[f'ABS(N(D{S_0}))>=0.005'],
                                                                font=Font(color='FF7030A0', bold=True)))
    ws.conditional_formatting.add(f'I{S_0}:J{S_1}', FormulaRule(formula=[f'LEFT($I{S_0},4)="首月手加"'],
                                                                fill=CF('FFFFC7CE'), font=Font(color='FF9C0006', bold=True)))
    rr = S_RES
    put(ws, f'A{rr}', '', fill=FILL_HDR)
    put(ws, f'B{rr}', '名字不在【门店月度收支统计】里的（也按门店算）')
    put(ws, f'C{rr}', None, fill=FILL_AUTO)
    put(ws, f'D{rr}', None, fill=FILL_AUTO)
    put(ws, f'E{rr}', f'={cut_sum("F", XR("C"), "1")}-SUM(E{S_0}:E{S_1})', fill=FILL_AUTO, fmt=MONEY, align=AR)
    put(ws, f'F{rr}', f'={cut_sum("G", XR("C"), "1")}-SUM(F{S_0}:F{S_1})', fill=FILL_AUTO, fmt=MONEY, align=AR)
    put(ws, f'G{rr}', f'=E{rr}-F{rr}', font=F_TXTB, fill=FILL_AUTO, fmt=MONEY, align=AR)
    put(ws, f'H{rr}', None, fill=FILL_AUTO)
    gx = f'{AUX}!$G${D0}:$G${D1}'
    merge(ws, f'I{rr}:J{rr}', f'=IF(AND(ABS(E{rr})<0.005,ABS(F{rr})<0.005),"",IF(COUNT({gx})=0,'
                              f'"门店超过 {S_1 - S_0 + 1} 家放不下，要请人加行",'
                              f'"数据录入第 "&MIN({gx})&" 行「"&INDEX(数据录入!$I:$I,MIN({gx}))&"」不在门店月度收支统计里（注意前后空格）：'
                              f'是门店就改成一样，不是门店就加到 ④"))', font=F_NOTE)
    ws.row_dimensions[rr].height = 44
    ws.conditional_formatting.add(f'E{rr}:G{rr}', FormulaRule(formula=[f'OR(ABS($E${rr})>=0.005,ABS($F${rr})>=0.005)'],
                                                              fill=CF('FFFFC7CE')))
    rr = S_TOT
    put(ws, f'A{rr}', '', fill=FILL_TOT)
    put(ws, f'B{rr}', '合计（＝ ① 的门店余额）', font=F_TXTB, fill=FILL_TOT)
    for col in 'CDEFGH':
        last = S_1 if col in 'CDH' else S_RES
        put(ws, f'{col}{rr}', f'=SUM({col}{S_0}:{col}{last})', font=F_TXTB, fill=FILL_TOT, fmt=MONEY, align=AR)
    merge(ws, f'I{rr}:J{rr}', '', fill=FILL_TOT)

    # ---------------- ⑥ 账户余额 ----------------
    section(ws, A_H, '⑥ 账户余额（账户＝【数据录入】第 5～18 行；G 列可以手填银行对账单上的余额）')
    heads(ws, A_H + 1, [('A', '序'), ('B', '账户'), ('C', '年初余额'), ('D', '收入'), ('E', '支出'), ('F', '账上余额'),
                        ('G', '银行对账单余额\n（手填，可空）'), ('H', '差\n（对账单－账上）'), ('I', '对账单余额\n（没填按账上）')], 40)
    acc = '数据录入!$C$5:$C$18'
    for i in range(A_1 - A_0 + 1):
        rr = A_0 + i
        b = f'$B{rr}'
        put(ws, f'A{rr}', f'=IF({b}="","",{i + 1})', fill=FILL_HDR, align=AC)
        put(ws, f'B{rr}', f'=IF({i + 1}>ROWS({acc}),"",INDEX({acc},{i + 1})&"")')
        put(ws, f'C{rr}', f'=IF({b}="","",N(INDEX(数据录入!$H$5:$H$18,{i + 1})))', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'D{rr}', f'=IF({b}="","",{cut_sum("F", DR("C"), b)})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{rr}', f'=IF({b}="","",{cut_sum("G", DR("C"), b)})', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{rr}', f'=IF({b}="","",ROUND(C{rr}+D{rr}-E{rr},2))', font=F_TXTB, fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{rr}', None, font=F_IN, fill=FILL_IN, fmt=MONEY, align=AR).protection = UNLOCK
        put(ws, f'H{rr}', f'=IF(OR({b}="",NOT(ISNUMBER(G{rr}))),"",ROUND(G{rr}-F{rr},2))', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'I{rr}', f'=IF({b}="","",IF(ISNUMBER(G{rr}),G{rr},F{rr}))', fill=FILL_AUTO, fmt=MONEY, align=AR)
        put(ws, f'J{rr}', f'=IF(AND({b}<>"",ISNUMBER(G{rr})),1,0)', font=Font(name=YH, size=8, color='FFFFFFFF'), border=None)
    dva = DataValidation(type='decimal', operator='between', formula1='-9999999999', formula2='9999999999',
                         allow_blank=True, showErrorMessage=True, errorTitle='只能填数字', error='填对账单上的余额')
    ws.add_data_validation(dva)
    dva.add(f'G{A_0}:G{A_1}')
    ws.conditional_formatting.add(f'H{A_0}:H{A_1}', FormulaRule(formula=[f'AND(H{A_0}<>"",ABS(N(H{A_0}))>=0.005)'],
                                                                fill=CF('FFFFC7CE')))
    ws.conditional_formatting.add(f'G{A_0}:G{A_1}', FormulaRule(formula=[f'$B{A_0}=""'], fill=CF('FFF2F2F2')))
    rr = A_TOT
    put(ws, f'A{rr}', '', fill=FILL_TOT)
    put(ws, f'B{rr}', '合计（账上余额＝ ① 的实存）', font=F_TXTB, fill=FILL_TOT)
    for col in 'CDEFI':
        put(ws, f'{col}{rr}', f'=ROUND(SUM({col}{A_0}:{col}{A_1}),2)', font=F_TXTB, fill=FILL_TOT, fmt=MONEY, align=AR)
    for col in 'GH':
        put(ws, f'{col}{rr}', f'=IF(SUM($J${A_0}:$J${A_1})=0,"",ROUND(SUM({col}{A_0}:{col}{A_1}),2))', font=F_TXTB,
            fill=FILL_TOT, fmt=MONEY, align=AR)
    put(ws, f'J{rr}', None, fill=FILL_TOT)
    merge(ws, f'A{A_TOT + 1}:J{A_TOT + 1}', 'G 列：截止日期那天银行 App / 对账单上的余额（现金填盘点数）。不填就按账上余额算；'
                                          '截止日期变了，这一列要跟着改。', font=F_NOTE, border=None)

    ws.column_dimensions['K'].hidden = True
    ws.column_dimensions['L'].hidden = True
    ws.freeze_panes = 'A6'
    ws.print_options.horizontalCentered = True
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = 9
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:2'
    ws.print_area = f'A1:J{A_TOT + 1}'
    for br in (X_H - 1, N_H - 1, S_H - 1, A_H - 1):
        ws.row_breaks.append(Break(id=br))
    # 保护（不设密码）：只有黄格子能填，防止把公式盖掉；要改格式：审阅 → 撤销工作表保护
    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.formatCells = False
    os.makedirs(TMP, exist_ok=True)
    p = os.path.join(TMP, 'main.xlsx')
    wb.save(p)
    return p


# ============================================================ 【门店余额核对辅助】（直接写 XML，用共享公式）
def aux_cells():
    """返回 {行: [(列, 公式或None, 文字或None, 共享信息)]}"""
    rows = {}

    def add(r, col, f=None, text=None, shared=None):
        rows.setdefault(r, []).append((col, f, text, shared))

    add(1, 'A', text=f'【门店余额核对】用的辅助表：不要改、不要删。A～G 列按行号跟【数据录入】一行对一行（第 {D0}～{D1} 行），'
                     '数据录入插行、删行也不会错位')
    for col, t in zip('ABCDEFG', ['计入日期', '账户：1 银行 2 空着 0 其他', '门店：1 本年门店 2 门店上年款 0 不算门店',
                                  '摘要含调拨', '金额格数', '所属月份：0 本年 1 空/错 2 上年 3 以后', '门店名不在统计表里（行号）']):
        add(3, col, text=t)
    for col, t in zip('HIJ', ['门店月度收支统计 B 列', '是门店且第一次出现', '累计第几个门店']):
        add(4, col, text=t)
    labels = ['截止日期（数字）', '最后一笔的计入日期', '年初银行余额（14 个账户）', '账本年份']
    vals = [f'{SH}!$L${R_CUT}', f'MAX($A${D0}:$A${D1})',
            'SUM(数据录入!$H$5:$H$18)', 'YEAR(数据录入!$B$5)']
    for i, (lab, f) in enumerate(zip(labels, vals), 1):
        add(i, 'L', text=lab)
        add(i, 'M', f=f)
    X_ = lambda c: f'INDEX(数据录入!${c}:${c},ROW())'
    B_, L_ = X_('B'), X_('L')
    y0, y1 = f'{MY}*100+1', f'{MY}*100+12'
    # 所属月份填错（文字、位数不对）不能报错：一律按 0（年初）算，③ 会数出来
    fa = (f'IFERROR(IF(AND(ISNUMBER({B_}),{B_}>=36526,{B_}<=2958465),INT({B_}),'
          f'IF(AND(ISNUMBER({L_}),{L_}>=190001,{L_}<=($M$4+1)*100+12,MOD({L_},100)>=1,MOD({L_},100)<=12),'
          f'DATE(INT({L_}/100),MOD({L_},100)+1,0),0)),0)')
    fb = f'IF({X_("C")}&""="",2,IF(COUNTIF(数据录入!$C$5:$C$18,{X_("C")})>0,1,0))'
    fc = (f'IF({X_("I")}&""="",0,IF(COUNTIF({SH}!{NSR},{X_("I")})>0,0,'
          f'IF(AND(ISNUMBER({L_}),{L_}>=190001,{L_}<$M$4*100+1),2,1)))')
    fd = f'IF(ISNUMBER(SEARCH("调拨",{X_("E")}&"")),1,0)'
    fe = f'ISNUMBER({X_("F")})+ISNUMBER({X_("G")})'
    ff = (f'IF(AND(ISNUMBER({L_}),{L_}>=190001),IF({L_}<{y0},2,IF({L_}>{y1},3,0)),1)')
    fg = f'IF(C{D0}<>1,"",IF(COUNTIF($H${M0}:$H${M1},{X_("I")}&"")=0,ROW(),""))'
    for col, f, si in (('A', fa, 0), ('B', fb, 1), ('C', fc, 2), ('D', fd, 3), ('E', fe, 4), ('F', ff, 5), ('G', fg, 9)):
        add(D0, col, f=f, shared=(si, f'{col}{D0}:{col}{D1}'))
        for r in range(D0 + 1, D1 + 1):
            add(r, col, shared=(si, None))
    # 门店名单（从门店月度收支统计 B 列取不重复的门店名，也按行号取）
    fh = f'IFERROR(INDEX(门店月度收支统计!$B:$B,ROW())&"","")'
    fi = f'IF(H{M0}="",0,IF(COUNTIF({SH}!{NSR},H{M0})>0,0,IF(COUNTIF($H${M0}:H{M0},H{M0})=1,1,0)))'
    add(M0, 'H', f=fh, shared=(6, f'H{M0}:H{M1}'))
    add(M0, 'I', f=fi, shared=(7, f'I{M0}:I{M1}'))
    add(M0, 'J', f=f'I{M0}')
    add(M0 + 1, 'J', f=f'J{M0}+I{M0 + 1}', shared=(8, f'J{M0 + 1}:J{M1}'))
    for r in range(M0 + 1, M1 + 1):
        add(r, 'H', shared=(6, None))
        add(r, 'I', shared=(7, None))
        if r > M0 + 1:
            add(r, 'J', shared=(8, None))
    return rows


def col_num(c):
    n = 0
    for ch in c:
        n = n * 26 + ord(ch) - 64
    return n


def esc(s):
    return (s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;'))


def vxml(v):
    """缓存值 → (t 属性, <v> 内容)"""
    if v is None:
        return None, None
    if isinstance(v, bool):
        return 'b', '1' if v else '0'
    if isinstance(v, (int, float)):
        return None, repr(float(v)) if isinstance(v, float) else str(v)
    if hasattr(v, 'toordinal'):          # 日期 → Excel 序列号
        import datetime as dt
        base = dt.datetime(1899, 12, 30)
        if not isinstance(v, dt.datetime):
            v = dt.datetime(v.year, v.month, v.day)
        return None, repr((v - base).total_seconds() / 86400.0).rstrip('0').rstrip('.')
    s = str(v)
    if s.startswith('#'):
        return 'e', s
    return 'str', s


def aux_xml(cache, hdr_style):
    rows = aux_cells()
    out = ['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
           f'<worksheet xmlns="{NSX}" xmlns:r="{RNS}">',
           '<sheetPr><tabColor rgb="FF7F7F7F"/></sheetPr>',
           f'<dimension ref="A1:M{D1}"/>',
           '<sheetViews><sheetView workbookViewId="0"><pane ySplit="4" topLeftCell="A5" activePane="bottomLeft" state="frozen"/>'
           '</sheetView></sheetViews>',
           '<sheetFormatPr defaultRowHeight="15"/>',
           '<cols><col min="1" max="7" width="13" customWidth="1"/>'
           '<col min="8" max="8" width="22" customWidth="1"/><col min="9" max="10" width="12" customWidth="1"/>'
           '<col min="11" max="11" width="3" customWidth="1"/><col min="12" max="12" width="24" customWidth="1"/>'
           '<col min="13" max="13" width="16" customWidth="1"/></cols>',
           '<sheetData>']
    for r in sorted(rows):
        cells = sorted(rows[r], key=lambda x: col_num(x[0]))
        parts = [f'<row r="{r}">']
        for col, f, text, shared in cells:
            ref = f'{col}{r}'
            if text is not None:
                st = f' s="{hdr_style}"' if r in (1, 3, 4) or col == 'L' else ''
                parts.append(f'<c r="{ref}"{st} t="inlineStr"><is><t>{esc(text)}</t></is></c>')
                continue
            t, v = vxml(cache.get(ref))
            ta = f' t="{t}"' if t else ''
            if shared is not None:
                si, rng = shared
                fx = (f'<f t="shared" ref="{rng}" si="{si}">{esc(f)}</f>' if rng else f'<f t="shared" si="{si}"/>')
            else:
                fx = f'<f>{esc(f)}</f>'
            vx = f'<v>{esc(v)}</v>' if v is not None else ''
            parts.append(f'<c r="{ref}"{ta}>{fx}{vx}</c>')
        parts.append('</row>')
        out.append(''.join(parts))
    out.append('</sheetData><sheetProtection sheet="1" objects="1" scenarios="1"/>'
               '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/></worksheet>')
    return '\n'.join(out).encode('utf-8')


# ============================================================ 样式合并
def merge_styles(orig_xml, tmp_xml):
    """把临时工作簿的样式接到原文件 styles.xml 后面，返回 (新 styles.xml, xf 映射, dxf 映射, 辅助表表头样式号)"""
    o = etree.fromstring(orig_xml)
    t = etree.fromstring(tmp_xml)

    def coll(root, tag):
        return root.find(NS + tag)

    # numFmts
    onf = coll(o, 'numFmts')
    if onf is None:
        onf = etree.Element(NS + 'numFmts', count='0')
        o.insert(0, onf)
    used = {int(n.get('numFmtId')) for n in onf}
    nxt = max(used | {200}) + 1
    nf_map = {}
    tnf = coll(t, 'numFmts')
    if tnf is not None:
        for n in tnf:
            old = int(n.get('numFmtId'))
            code = n.get('formatCode')
            same = [int(x.get('numFmtId')) for x in onf if x.get('formatCode') == code]
            if same:
                nf_map[old] = same[0]
                continue
            e = etree.SubElement(onf, NS + 'numFmt', numFmtId=str(nxt), formatCode=code)
            nf_map[old] = nxt
            nxt += 1
    onf.set('count', str(len(onf)))

    offs = {}
    for tag in ('fonts', 'fills', 'borders'):
        oc, tc = coll(o, tag), coll(t, tag)
        offs[tag] = len(oc)
        for e in tc:
            oc.append(copy.deepcopy(e))
        oc.set('count', str(len(oc)))
    ox, tx = coll(o, 'cellXfs'), coll(t, 'cellXfs')
    xf0 = len(ox)
    xf_map = {}
    for i, e in enumerate(tx):
        e = copy.deepcopy(e)
        nf = int(e.get('numFmtId', '0'))
        e.set('numFmtId', str(nf_map.get(nf, nf)))
        e.set('fontId', str(int(e.get('fontId', '0')) + offs['fonts']))
        e.set('fillId', str(int(e.get('fillId', '0')) + offs['fills']))
        e.set('borderId', str(int(e.get('borderId', '0')) + offs['borders']))
        e.set('xfId', '0')
        ox.append(e)
        xf_map[i] = xf0 + i
    # 辅助表表头：微软雅黑粗体
    of = coll(o, 'fonts')
    fe = etree.SubElement(of, NS + 'font')
    etree.SubElement(fe, NS + 'b')
    etree.SubElement(fe, NS + 'sz', val='10')
    etree.SubElement(fe, NS + 'name', val='微软雅黑')
    etree.SubElement(fe, NS + 'charset', val='134')
    of.set('count', str(len(of)))
    etree.SubElement(ox, NS + 'xf', numFmtId='0', fontId=str(len(of) - 1), fillId='0', borderId='0', xfId='0',
                     applyFont='1')
    hdr_idx = len(ox) - 1
    ox.set('count', str(len(ox)))
    od, td = coll(o, 'dxfs'), coll(t, 'dxfs')
    dxf_map = {}
    if td is not None and len(td):
        if od is None:
            od = etree.Element(NS + 'dxfs', count='0')
            coll(o, 'cellStyles').addnext(od)
        d0 = len(od)
        for i, e in enumerate(td):
            od.append(copy.deepcopy(e))
            dxf_map[i] = d0 + i
        od.set('count', str(len(od)))
    xml = etree.tostring(o, xml_declaration=True, encoding='UTF-8', standalone=True)
    return xml, xf_map, dxf_map, hdr_idx


def main_xml(tmp_path, xf_map, dxf_map, cache):
    z = zipfile.ZipFile(tmp_path)
    sst = []
    if 'xl/sharedStrings.xml' in z.namelist():
        s = etree.fromstring(z.read('xl/sharedStrings.xml'))
        for si in s.findall(NS + 'si'):
            sst.append(''.join(x.text or '' for x in si.iter(NS + 't')))
    root = etree.fromstring(z.read('xl/worksheets/sheet1.xml'))
    for sv in root.iter(NS + 'sheetView'):
        if 'tabSelected' in sv.attrib:
            del sv.attrib['tabSelected']
    for c in root.iter(NS + 'c'):
        s = c.get('s')
        c.set('s', str(xf_map.get(int(s) if s else 0, xf_map[0])))
        if c.get('t') == 's':
            v = c.find(NS + 'v')
            text = sst[int(v.text)]
            c.remove(v)
            c.set('t', 'inlineStr')
            is_ = etree.SubElement(c, NS + 'is')
            tt = etree.SubElement(is_, NS + 't')
            tt.text = text
            if text != text.strip():
                tt.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        f = c.find(NS + 'f')
        if f is not None:
            v = c.find(NS + 'v')
            if v is not None:
                c.remove(v)
            if 't' in c.attrib:
                del c.attrib['t']
            t, val = vxml(cache.get(c.get('r')))
            if val is not None:
                if t:
                    c.set('t', t)
                ve = etree.SubElement(c, NS + 'v')
                ve.text = val
    for rule in root.iter(NS + 'cfRule'):
        if rule.get('dxfId') is not None:
            rule.set('dxfId', str(dxf_map[int(rule.get('dxfId'))]))
    # 没写 s 的格子（openpyxl 默认样式）也要指到新样式 0
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True), \
        z.read('xl/styles.xml')


# ============================================================ 拼包
def assemble(src, dst, cache_main, cache_aux):
    tmp = build_main()
    zin = zipfile.ZipFile(src)
    names = zin.namelist()
    styles_new, xf_map, dxf_map, hdr_idx = merge_styles(zin.read('xl/styles.xml'), zipfile.ZipFile(tmp).read('xl/styles.xml'))
    sheet_main, _ = main_xml(tmp, xf_map, dxf_map, cache_main)
    sheet_aux = aux_xml(cache_aux, hdr_idx)

    wbx = zin.read('xl/workbook.xml').decode('utf-8')
    rels = zin.read('xl/_rels/workbook.xml.rels').decode('utf-8')
    ct = zin.read('[Content_Types].xml').decode('utf-8')
    app = zin.read('docProps/app.xml').decode('utf-8')
    if f'name="{SH}"' in wbx:
        raise SystemExit('原表里已经有【门店余额核对】了：请用没加过的原表')
    sheet_ids = [int(x) for x in re.findall(r'<sheet [^>]*sheetId="(\d+)"', wbx)]
    rids = [int(x) for x in re.findall(r'Id="rId(\d+)"', rels)]
    parts = [int(x) for x in re.findall(r'worksheets/sheet(\d+)\.xml', rels)]
    sid1, sid2 = max(sheet_ids) + 1, max(sheet_ids) + 2
    rid1, rid2 = f'rId{max(rids) + 1}', f'rId{max(rids) + 2}'
    p1, p2 = f'sheet{max(parts) + 1}.xml', f'sheet{max(parts) + 2}.xml'
    idx = len(re.findall(r'<sheet [^>]*/>', wbx))          # 新表在工作表清单里的位置（从 0 数）
    dn = (f'<definedName name="_xlnm.Print_Titles" localSheetId="{idx}">{SH}!$1:$2</definedName>'
          f'<definedName name="_xlnm.Print_Area" localSheetId="{idx}">{SH}!$A$1:$J${A_TOT + 1}</definedName>')
    if '</definedNames>' in wbx:
        wbx = wbx.replace('</definedNames>', dn + '</definedNames>', 1)
    else:
        wbx = wbx.replace('</sheets>', '</sheets><definedNames>' + dn + '</definedNames>', 1)
    wbx = wbx.replace('</sheets>', f'<sheet name="{SH}" sheetId="{sid1}" r:id="{rid1}"/>'
                                   f'<sheet name="{AUX}" sheetId="{sid2}" state="hidden" r:id="{rid2}"/></sheets>', 1)
    ws_t = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet'
    rels = rels.replace('</Relationships>', f'<Relationship Id="{rid1}" Type="{ws_t}" Target="worksheets/{p1}"/>'
                                            f'<Relationship Id="{rid2}" Type="{ws_t}" Target="worksheets/{p2}"/></Relationships>', 1)
    ws_ct = 'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'
    ct = ct.replace('</Types>', f'<Override PartName="/xl/worksheets/{p1}" ContentType="{ws_ct}"/>'
                                f'<Override PartName="/xl/worksheets/{p2}" ContentType="{ws_ct}"/></Types>', 1)
    m = re.search(r'<vt:lpstr>工作表</vt:lpstr></vt:variant><vt:variant><vt:i4>(\d+)</vt:i4>', app)
    if m:
        n = int(m.group(1))
        app = app.replace(m.group(0), m.group(0).replace(f'<vt:i4>{n}</vt:i4>', f'<vt:i4>{n + 2}</vt:i4>'), 1)
        m2 = re.search(r'<TitlesOfParts><vt:vector size="(\d+)" baseType="lpstr">', app)
        if m2:
            k = int(m2.group(1))
            app = app.replace(m2.group(0), m2.group(0).replace(f'size="{k}"', f'size="{k + 2}"'), 1)
            app = app.replace('</vt:vector></TitlesOfParts>',
                              f'<vt:lpstr>{SH}</vt:lpstr><vt:lpstr>{AUX}</vt:lpstr></vt:vector></TitlesOfParts>', 1)
    repl = {'xl/workbook.xml': wbx.encode('utf-8'), 'xl/_rels/workbook.xml.rels': rels.encode('utf-8'),
            '[Content_Types].xml': ct.encode('utf-8'), 'docProps/app.xml': app.encode('utf-8'),
            'xl/styles.xml': styles_new}
    zout = zipfile.ZipFile(dst + '.part', 'w', zipfile.ZIP_DEFLATED)
    for info in zin.infolist():
        data = repl.get(info.filename, zin.read(info.filename))
        zi = zipfile.ZipInfo(info.filename, date_time=info.date_time)
        zi.compress_type = info.compress_type          # 目录项原来是不压缩的，照原样
        zi.create_system = info.create_system
        zi.external_attr = info.external_attr
        zout.writestr(zi, data)
    zout.writestr(f'xl/worksheets/{p1}', sheet_main, zipfile.ZIP_DEFLATED)
    zout.writestr(f'xl/worksheets/{p2}', sheet_aux, zipfile.ZIP_DEFLATED)
    zout.close()
    os.replace(dst + '.part', dst)
    return p1, p2


def recalc_values(path):
    """副本上用 LibreOffice 整本重算，读回新加两张表的数"""
    rc = (glob.glob('/root/.claude/skills/synced/*/xlsx/scripts/recalc.py') +
          glob.glob('/mnt/skills/public/xlsx/scripts/recalc.py'))[0]
    cp = os.path.join(TMP, 'calc.xlsx')
    shutil.copy(path, cp)
    r = subprocess.run([sys.executable, rc, cp, '1800', '--force'], capture_output=True, text=True)
    print('   LibreOffice：', r.stdout.strip()[:400])
    wb = load_workbook(cp, data_only=True, read_only=True)
    out = []
    for nm in (SH, AUX):
        d = {}
        for row in wb[nm].iter_rows():
            for c in row:
                if c.value is not None:
                    d[c.coordinate] = c.value
        out.append(d)
    return out


if __name__ == '__main__':
    os.makedirs(TMP, exist_ok=True)
    print('① 拼新表（还没缓存值）')
    assemble(SRC, OUT, {}, {})
    print('② LibreOffice 副本上算数')
    cm, ca = recalc_values(OUT)
    print(f'   新表 {len(cm)} 格有数，辅助表 {len(ca)} 格')
    print('③ 把算好的数写进成品当缓存')
    assemble(SRC, OUT, cm, ca)
    print('已保存：', OUT, f'{os.path.getsize(OUT) / 1e6:.1f} MB')
