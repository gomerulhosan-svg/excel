# 两个通用核对管不到的场景：删账户期初行（AcctDel）、复制核对表（CopySh）
import os
import sys
from openpyxl import load_workbook
HERE = os.path.dirname(os.path.abspath(__file__))
base = load_workbook(os.path.join(HERE, '..', '..', '多帐户收支登记表_加门店余额核对.xlsx'), data_only=True, read_only=True)['门店余额核对']
bv = {c.coordinate: c.value for row in base.iter_rows(min_row=1, max_row=180, max_col=10) for c in row if hasattr(c, 'coordinate')}
bad = []
def chk(name, got, want, tol=0.005):
    ok = (abs(got - want) <= tol) if isinstance(want, (int, float)) and isinstance(got, (int, float)) else got == want
    print(('  ok ' if ok else '  ✗  ') + name, got, want)
    if not ok: bad.append(name)

print('== AcctDel：删掉第 18 行（五优43店邮政 期初），第一笔流水挤到第 18 行')
wb = load_workbook('AcctDel.xlsx', data_only=True, read_only=True)
ws = wb['门店余额核对']; v = {c.coordinate: c.value for row in ws.iter_rows(min_row=1, max_row=180, max_col=10) for c in row if hasattr(c, 'coordinate')}
d = wb['数据录入']; dr = list(d.iter_rows(min_row=1, max_row=19, max_col=12, values_only=True))
aux = wb['门店余额核对辅助']; ar = {c.coordinate: c.value for row in aux.iter_rows(min_row=17, max_row=19, max_col=7) for c in row if hasattr(c, 'coordinate')}
print('   数据录入第 18 行：', dr[17][1], dr[17][2], dr[17][4], dr[17][5], dr[17][8], dr[17][11])
for ref, lab in (('C9', '门店余额'), ('C10', '收款'), ('C11', '付款')):
    chk(f'{lab} 不变（第一笔流水还在）', v[ref], bv[ref])
chk('应存＝H3', v['C16'], d['H3'].value)
chk('辅助第 18 行就是第一笔流水（门店上年款 C=2）', ar['C18'], 2)
chk('⑥ 邮政个人 收入不变', v['D165'], bv['D165'])
chk('自查', v['G47'], '✓ 一致')
chk('③ H3 说明', v['G44'], '＝ ① 的应存（不是银行余额）')

print('== CopySh：复制核对表，副本 B4 填 2026-06-30')
wb = load_workbook('CopySh.xlsx', data_only=True, read_only=True)
cp = wb['核对副本']; c = {x.coordinate: x.value for row in cp.iter_rows(min_row=1, max_row=40, max_col=10) for x in row if hasattr(x, 'coordinate')}
orig = wb['门店余额核对']; o = {x.coordinate: x.value for row in orig.iter_rows(min_row=1, max_row=40, max_col=10) for x in row if hasattr(x, 'coordinate')}
chk('副本 H4', c['H4'], '截至 2026-06-30')
chk('副本 应存＝② 6 月', c['C16'], c['G28'])
chk('副本 实存＝② 6 月', c['C17'], c['H28'])
chk('副本 差额＝② 6 月', c['C18'], c['I28'])
chk('副本 差额', c['C18'], -12559.16)
chk('原表 差额不受影响', o['C18'], -131037.0)
chk('原表 H4 不受影响', o['H4'], bv['H4'])
print('不一致：', bad if bad else '无')
sys.exit(1 if bad else 0)
