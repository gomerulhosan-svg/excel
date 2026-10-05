# -*- coding: utf-8 -*-
"""独立核对：不看新表的公式，直接从原表的【数据录入】【门店月度收支统计】用 Python 重算一遍，跟成品里的数逐格比；
再核对原文件的每个部件都原样保留（只有 workbook / rels / Content_Types / app / styles 五个文件加了东西）。

跑法：python3 verify.py [成品] [原表]
     python3 verify.py 改过的成品 --values   只核对数（场景测试用：数据就从这本里读，截止日期、对账单、④ 名字也照这本里填的算）
"""
import collections
from decimal import Decimal, ROUND_HALF_UP
import datetime as dt
import os
import re
import sys
import zipfile

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from build import LAYOUT  # noqa: E402  只借行号，公式一个都不借

ARGS = [a for a in sys.argv[1:] if not a.startswith('--')]
VALUES_ONLY = '--values' in sys.argv
OUT = ARGS[0] if ARGS else os.path.join(HERE, '..', '多帐户收支登记表_加门店余额核对.xlsx')
SRC = ARGS[1] if len(ARGS) > 1 else (OUT if VALUES_ONLY else os.path.join(HERE, '..', '参考', '原表_数据录入_2026.9.27.xlsx'))
SH, AUX = '门店余额核对', '门店余额核对辅助'
R1, X, M_0 = LAYOUT['R1'], LAYOUT['X'], LAYOUT['M_0']
N_0, N_1, N_PREV, N_BL, N_TOT = (LAYOUT[k] for k in ('N_0', 'N_1', 'N_PREV', 'N_BL', 'N_TOT'))
S_0, S_1, S_RES, S_TOT = (LAYOUT[k] for k in ('S_0', 'S_1', 'S_RES', 'S_TOT'))
A_0, A_1, A_TOT = (LAYOUT[k] for k in ('A_0', 'A_1', 'A_TOT'))
BAD = []
N_OK = [0]


def chk(name, got, want, tol=0.005):
    if isinstance(got, dt.datetime) and isnum(want):          # LibreOffice 存盘时给 DATE() 的结果自动套了日期格式
        got = (got - dt.datetime(1899, 12, 30)).days
    if isinstance(want, (int, float)) and not isinstance(want, bool):
        ok = isinstance(got, (int, float)) and not isinstance(got, bool) and abs(got - want) <= tol
    else:
        ok = (got if got is not None else '') == (want if want is not None else '')
    if ok:
        N_OK[0] += 1
    else:
        BAD.append((name, got, want))


def r2(x):
    return float(Decimal(repr(float(x))).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def isnum(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def num(v):
    return v if isnum(v) else 0


def txt(v):
    return '' if v is None else (str(int(v)) if isinstance(v, float) and v == int(v) else str(v))


def serial(d):
    return (d - dt.datetime(1899, 12, 30)).days


def money(x):
    return f'{abs(x):,.2f}'


# ---------------- 成品 ----------------
wb = load_workbook(OUT, data_only=True)
ws, wa = wb[SH], wb[AUX]
v = lambda ref: ws[ref].value
NS_ROWS = [(rr, txt(v(f'B{rr}'))) for rr in range(N_0, N_1 + 1)]
NONSTORE = [n for _, n in NS_ROWS if n]                     # ④ B 列填的名字（默认 9 个）
CUT = v('B4') if v('B4') is not None else 2958465
if isinstance(CUT, dt.datetime):
    CUT = serial(CUT)
STMT = {i: v(f'G{A_0 + i}') for i in range(A_1 - A_0 + 1)}  # ⑥ 手填的对账单余额

# ---------------- 原表数据 ----------------
src = load_workbook(SRC, data_only=True, read_only=True)
de = [list(r) for r in src['数据录入'].iter_rows(min_row=1, max_row=30191, max_col=15, values_only=True)]
row = lambda r: de[r - 1] if r - 1 < len(de) else [None] * 15
accts = [row(r)[2] for r in range(5, 19)]
bank0 = sum(num(row(r)[7]) for r in range(5, 19))
book_y = row(5)[1].year
H3 = row(3)[7]
ms = [list(r) for r in src['门店月度收支统计'].iter_rows(min_row=1, max_row=1004, max_col=11, values_only=True)]
ms_year = ms[1][1]
msr = lambda r: ms[r - 1] if r - 1 < len(ms) else [None] * 11


def eff(x):
    b, L = x[1], x[11]
    if isinstance(b, dt.datetime):
        return serial(b)
    if isnum(b) and 36526 <= b <= 2958465:
        return int(b)
    if isnum(L) and L >= 190001 and 1 <= L % 100 <= 12:
        y, m = int(L // 100), int(L % 100)
        nxt = dt.datetime(y + (m == 12), m % 12 + 1, 1)
        return serial(nxt - dt.timedelta(days=1))
    return 0


def mtype(L):
    if isnum(L) and L >= 190001:
        return 2 if L < ms_year * 100 + 1 else (3 if L > ms_year * 100 + 12 else 0)
    return 1


flows = []
for r in range(19, 30192):
    x = list(row(r))
    if isinstance(x[11], dt.datetime):                        # 插的行从上一行带了日期格式：202609 会读成日期，还原成数
        x[11] = (x[11] - dt.datetime(1899, 12, 30)).days
    C, I, L = txt(x[2]), txt(x[8]), x[11]
    bank = 2 if C == '' else (1 if C in accts else 0)
    if I == '' or I in NONSTORE:
        store = 0
    elif isnum(L) and L >= 190001 and L < book_y * 100 + 1:
        store = 2                                     # 门店上年的款
    else:
        store = 1
    flows.append(dict(r=r, eff=eff(x), bank=bank, store=store, tiao=1 if '调拨' in txt(x[4]) else 0,
                      n=(txt(x[5]) != '') + (txt(x[6]) != ''), mt=mtype(L),
                      F=num(x[5]), G=num(x[6]), C=C, I=I, L=L))

# 门店名单（门店月度收支统计 B 列，不重复，去掉不算门店的名字）
names, seen = [], set()
for r in range(5, 1005):
    b = txt(msr(r)[1])
    if b == '' or b in NONSTORE or b in seen:
        continue
    seen.add(b)
    names.append(b)
msG, msD, msE, msN = collections.defaultdict(float), collections.defaultdict(float), collections.defaultdict(float), \
    collections.Counter()
for r in range(5, 1005):
    x = msr(r)
    b = txt(x[1])
    if b:
        msG[b] += num(x[6]); msD[b] += num(x[3]); msE[b] += num(x[4]); msN[b] += 1


def agg(cut, pred):
    f = g = 0.0
    n = 0
    for x in flows:
        if x['eff'] <= cut and pred(x):
            f += x['F']; g += x['G']; n += x['n']
    return f, g, n


def store_tables(cut):
    out = []
    for nm in names:
        K = sum(x['F'] - x['G'] for x in flows if x['I'] == nm and x['mt'] == 0)
        f, g, _ = agg(cut, lambda x, nm=nm: x['I'] == nm and x['store'] == 1)
        miss = sum(x['F'] - x['G'] for x in flows if x['I'] == nm and x['store'] == 1 and x['mt'] != 0)
        adj = r2(msD[nm] - msE[nm] - K)
        out.append(dict(name=nm, open=msG[nm], adj=adj, F=f, G=g, end=msG[nm] + adj + f - g, miss=r2(miss)))
    return out


last = max(x['eff'] for x in flows)

st = store_tables(CUT)
open_sum = sum(s['open'] for s in st)
adj_sum = sum(s['adj'] for s in st)
sf, sg, _ = agg(CUT, lambda x: x['store'] == 1)
nf, ng, _ = agg(CUT, lambda x: x['store'] != 1)
bf, bg, _ = agg(CUT, lambda x: x['bank'] == 1)
co_a = r2(sum(msG[n] for n in NONSTORE))
co_b = r2(bank0 - open_sum - co_a)
co_c = -r2(adj_sum)
co = r2(co_a + co_b + co_c)
store_bal = open_sum + adj_sum + sf - sg
should = r2(store_bal + nf - ng + co)
bank = r2(bank0 + bf - bg)
diff = r2(bank - should)
day = lambda n: f'{(dt.datetime(1899, 12, 30) + dt.timedelta(days=n)):%Y-%m-%d}'
ALL = CUT == 2958465
p = '[全部] ' if ALL else f'[截至 {day(CUT)}] '
chk(p + '现在算的是', v('H4'), f'全部（最后一笔 {day(last)}）' if ALL else f'截至 {day(CUT)}')
chk(p + '① 门店余额', v(f'C{R1["store"]}'), store_bal)
chk(p + '① 收款', v(f'C{R1["recv"]}'), nf)
chk(p + '① 付款', v(f'C{R1["pay"]}'), ng)
chk(p + '① 非门店年初数', v(f'C{R1["co"]}'), co)
chk(p + '① 4a 配送公司年初', v(f'C{R1["co_a"]}'), co_a)
chk(p + '① 4b 年初没对平', v(f'C{R1["co_b"]}'), co_b)
chk(p + '① 4c 减首月手加', v(f'C{R1["co_c"]}'), co_c)
chk(p + '① 应存', v(f'C{R1["should"]}'), should)
chk(p + '① 实存', v(f'C{R1["bank"]}'), bank)
chk(p + '① 差额', v(f'C{R1["diff"]}'), diff)
filled = {i: g for i, g in STMT.items() if isnum(g) and i < len(accts)}
acc_end = {i: r2(num(row(5 + i)[7]) + agg(CUT, lambda x, a=a: x['C'] == a)[0] - agg(CUT, lambda x, a=a: x['C'] == a)[1])
           for i, a in enumerate(accts)}
stmt = r2(sum(filled.get(i, acc_end[i]) for i in acc_end)) if filled else ''
chk(p + '① 对账单', v(f'C{R1["stmt"]}'), stmt)
if stmt == '':
    want = '⑥ G 列还没填银行对账单余额（填了就能跟银行的实际余额对）'
elif abs(stmt - bank) < 0.005:
    want = '✓ 对账单跟账上余额一致'
else:
    want = f'✗ 对账单比账上{"少" if stmt < bank else "多"} {money(stmt - bank)} 元：看 ⑥ 是哪个账户'
chk(p + '① 对账单提示', v(f'D{R1["stmt"]}'), want)
chk(p + '① 差额提示', v(f'D{R1["diff"]}'),
    '✓ 账面对上了：门店余额＋收款－付款＋非门店年初数＝14 个账户的账上余额（跟银行实际余额核对：在 ⑥ 填对账单余额）'
    if abs(diff) < 0.005 else f'✗ 实存比应存{"少" if diff < 0 else "多"} {money(diff)} 元：看下面 ③ 是哪些记录造成的')
chk(p + '① 应存说明里的 H3', v(f'D{R1["should"]}'),
    f'＝【数据录入】最上面 H3 的「账户余额」{money(H3)}（没筛选时）。那个数把不走银行的记录也算进去了，所以不等于银行余额；银行余额看下一行'
    if ALL else '')
if ALL:
    chk(p + '恒等式：应存＝H3（没筛选、没填截止日期时）', should, r2(H3))

# ③ 没走银行的记录
tiao = agg(CUT, lambda x: x['bank'] == 2 and x['tiao'] == 1)
blank = agg(CUT, lambda x: x['bank'] == 2)
zd = agg(CUT, lambda x: x['C'] == '货款直抵费用')
other = agg(CUT, lambda x: x['bank'] == 0)
exp = dict(a=tiao, b=tuple(blank[i] - tiao[i] for i in range(3)), c=zd, d=tuple(other[i] - zd[i] for i in range(3)))
for k, (f, g, n) in exp.items():
    chk(p + f'③{k} 笔数', v(f'C{X[k]}'), n, 0)
    chk(p + f'③{k} 收入', v(f'D{X[k]}'), f)
    chk(p + f'③{k} 支出', v(f'E{X[k]}'), g)
    chk(p + f'③{k} 净额', v(f'F{X[k]}'), r2(f - g))
nb_net = sum(r2(f - g) for f, g, _ in exp.values())
chk(p + '③ 小计净额', v(f'F{X["sub"]}'), nb_net)
chk(p + '③ 对差额的影响', v(f'F{X["eff"]}'), r2(-nb_net))
chk(p + '恒等式：差额＝－没走银行的净额', diff, r2(-nb_net))
chk(p + '③ 跟①一致的提示', v(f'G{X["eff"]}'), '✓ 跟 ① 的差额一致')
chk(p + '③ H3', v(f'F{X["h3"]}'), H3 if ALL else '')
chk(p + '③ H3 说明', v(f'G{X["h3"]}'), '＝ ① 的应存（不是银行余额）' if ALL else '（填了截止日期就不比）')
inner = agg(CUT, lambda x: x['I'] == '内部转款')
chk(p + '③ 内部转款净额', v(f'F{X["inner"]}'), r2(inner[0] - inner[1]))
chk(p + '③ 没日期笔数', v(f'C{X["nodate"]}'), sum(x['n'] for x in flows if x['eff'] == 0), 0)
chk(p + '自查', v(f'G{X["chk"]}'), '✓ 一致')

# ④
for rr, n in NS_ROWS:
    if n:
        f, g, _ = agg(CUT, lambda x, n=n: x['I'] == n)
        bf_, bg_, _ = agg(CUT, lambda x, n=n: x['I'] == n and x['bank'] == 1)
        for col, want in zip('CDEFGL', (f, g, f - g, bf_, bg_, msG[n])):
            chk(p + f'④ {n} {col}', v(f'{col}{rr}'), want)
        chk(p + f'④ {n} 重复', v(f'K{rr}'), 1, 0)
    else:
        for col in 'BCDEFG':
            chk(p + f'④ 空行 {col}{rr}', v(f'{col}{rr}'), '')
        chk(p + f'④ 空行 L{rr}', v(f'L{rr}'), 0, 0)
prev = agg(CUT, lambda x: x['store'] == 2)
prevb = agg(CUT, lambda x: x['store'] == 2 and x['bank'] == 1)
for col, want in zip('CDEFG', (prev[0], prev[1], prev[0] - prev[1], prevb[0], prevb[1])):
    chk(p + f'④ 门店上年的款 {col}', v(f'{col}{N_PREV}'), want)
bl = agg(CUT, lambda x: x['I'] == '')
blb = agg(CUT, lambda x: x['I'] == '' and x['bank'] == 1)
for col, want in zip('CDEFG', (bl[0], bl[1], bl[0] - bl[1], blb[0], blb[1])):
    chk(p + f'④ 没填门店 {col}', v(f'{col}{N_BL}'), want)
nb = agg(CUT, lambda x: x['store'] != 1 and x['bank'] == 1)
for col, want in zip('CDEFG', (nf, ng, nf - ng, nb[0], nb[1])):
    chk(p + f'④ 合计 {col}', v(f'{col}{N_TOT}'), want)

# ⑤
adj_cnt = collections.Counter(s['adj'] for s in st)
for i, s in enumerate(st):
    rr = S_0 + i
    nm = s['name']
    chk(p + f'⑤ {nm} 序', v(f'A{rr}'), i + 1, 0)
    chk(p + f'⑤ {nm} 名字', v(f'B{rr}'), nm)
    for col, k in zip('CDEFGH', ('open', 'adj', 'F', 'G', 'end', 'miss')):
        chk(p + f'⑤ {nm} {col}', v(f'{col}{rr}'), s[k])
    if abs(s['adj']) >= 0.005 and adj_cnt[s['adj']] > 1:
        hint = '首月手加的数跟另一家一模一样，核对是不是复制错了'
    elif msN[nm] < 12:
        hint = '已关的店：余额是门店月度收支统计里单列的一行'
    elif abs(s['miss']) >= 0.005:
        hint = '有记录没填所属月份，门店月度收支统计没算到'
    else:
        hint = ''
    chk(p + f'⑤ {nm} 提示', v(f'I{rr}'), hint)
for rr in range(S_0 + len(st), S_1 + 1):
    for col in 'ABCDEFGHI':
        chk(p + f'⑤ 空行 {col}{rr}', v(f'{col}{rr}'), '')
unk = agg(CUT, lambda x: x['store'] == 1 and x['I'] not in seen)
chk(p + '⑤ 名字不在门店月度里的 收入', v(f'E{S_RES}'), unk[0])
chk(p + '⑤ 名字不在门店月度里的 支出', v(f'F{S_RES}'), unk[1])
for col, want in zip('CDEFGH', (open_sum, adj_sum, sf, sg, store_bal, sum(s['miss'] for s in st))):
    chk(p + f'⑤ 合计 {col}', v(f'{col}{S_TOT}'), want)

# ⑥
for i in range(A_1 - A_0 + 1):
    rr = A_0 + i
    if i < len(accts):
        a = accts[i]
        f, g, _ = agg(CUT, lambda x, a=a: x['C'] == a)
        o = num(row(5 + i)[7])
        e = r2(o + f - g)
        gi = filled.get(i)
        for col, want in zip('BCDEFHIJ', (a, o, f, g, e, '' if gi is None else r2(gi - e), e if gi is None else gi,
                                         0 if gi is None else 1)):
            chk(p + f'⑥ {a} {col}', v(f'{col}{rr}'), want)
    else:
        for col in 'ABCDEFHI':
            chk(p + f'⑥ 空行 {col}{rr}', v(f'{col}{rr}'), '')
        chk(p + f'⑥ 空行 J{rr}', v(f'J{rr}'), 0, 0)
chk(p + '⑥ 合计 年初', v(f'C{A_TOT}'), r2(bank0))
chk(p + '⑥ 合计＝实存', v(f'F{A_TOT}'), bank)
chk(p + '⑥ 合计 对账单', v(f'G{A_TOT}'), r2(sum(filled.values())) if filled else '')
chk(p + '⑥ 合计 差', v(f'H{A_TOT}'), r2(sum(r2(g - acc_end[i]) for i, g in filled.items())) if filled else '')
chk(p + '⑥ 合计 对账单（没填按账上）', v(f'I{A_TOT}'), stmt if filled else bank)
print('全部：门店', round(store_bal, 2), '收款', round(nf, 2), '付款', round(ng, 2), '非门店年初', co,
      '应存', should, '实存', bank, '差额', diff)

# ② 按月
OPEN_ADJ = open_sum + adj_sum
for k in range(1, 13):
    rr = M_0 + k - 1
    me = serial(dt.datetime(book_y + (k == 12), k % 12 + 1, 1) - dt.timedelta(days=1))
    if serial(dt.datetime(book_y, k, 1)) > last:
        for col in 'CDEFGHIJ':
            chk(f'② {k}月 空着 {col}', v(f'{col}{rr}'), '')
        continue
    sf_, sg_, _ = agg(me, lambda x: x['store'] == 1)
    nf_, ng_, _ = agg(me, lambda x: x['store'] != 1)
    bf_, bg_, _ = agg(me, lambda x: x['bank'] == 1)
    c, d, e = r2(OPEN_ADJ + sf_ - sg_), r2(nf_), r2(ng_)
    g = r2(c + d - e + co)
    h = r2(bank0 + bf_ - bg_)
    i_ = r2(h - g)
    for col, want in zip('CDEFGHI', (c, d, e, co, g, h, i_)):
        chk(f'② {k}月 {col}', v(f'{col}{rr}'), want)
    chk(f'② {k}月 标记', v(f'J{rr}'), '✓' if abs(i_) < 0.005 else ('✗ 少' if i_ < 0 else '✗ 多'))
    # 每个月的差额＝－到月底为止没走银行的净额
    nbm = agg(me, lambda x: x['bank'] != 1)
    chk(f'② {k}月 恒等式', i_, r2(-(nbm[0] - nbm[1])))

# 辅助表逐行
for x in flows:
    r = x['r']
    for col, key in zip('ABCDEF', ('eff', 'bank', 'store', 'tiao', 'n', 'mt')):
        chk(f'辅助 {col}{r}', wa[f'{col}{r}'].value, x[key], 0)
cum = 0
first = set()
for r in range(5, 1005):
    b = txt(msr(r)[1])
    chk(f'辅助 H{r}', wa[f'H{r}'].value, b)
    flag = 1 if (b and b not in NONSTORE and b not in first) else 0
    if b:
        first.add(b)
    cum += flag
    chk(f'辅助 I{r}', wa[f'I{r}'].value, flag, 0)
    chk(f'辅助 J{r}', wa[f'J{r}'].value, cum, 0)
chk('辅助 M1 截止日期', wa['M1'].value, CUT, 0)
chk('辅助 A～F 只到第 30191 行', wa.max_row, 30191, 0)
chk('辅助 M2 最后一笔', wa['M2'].value, last, 0)
chk('辅助 M3 年初银行余额', wa['M3'].value, bank0)
chk('辅助 M4 账本年份', wa['M4'].value, book_y, 0)

if VALUES_ONLY:
    print(f'核对 {N_OK[0] + len(BAD)} 项，一致 {N_OK[0]} 项，不一致 {len(BAD)} 项（只核对数）')
    for b in BAD[:40]:
        print('  ✗', b)
    sys.exit(1 if BAD else 0)

# ---------------- 原文件部件原样保留 ----------------
zs, zo = zipfile.ZipFile(SRC), zipfile.ZipFile(OUT)
changed = {'xl/workbook.xml', 'xl/_rels/workbook.xml.rels', '[Content_Types].xml', 'docProps/app.xml', 'xl/styles.xml'}
for n in zs.namelist():
    if n in changed:
        continue
    chk(f'部件原样 {n}', zo.read(n) == zs.read(n), True)
info_s = {i.filename: i for i in zs.infolist()}
for i in zo.infolist():
    if i.filename in info_s:
        o = info_s[i.filename]
        chk(f'压缩方式、属性原样 {i.filename}', (i.compress_type, i.create_system, i.external_attr),
            (o.compress_type, o.create_system, o.external_attr))
new_parts = sorted(set(zo.namelist()) - set(zs.namelist()))
chk('只多了两张表', len(new_parts), 2, 0)
wo, wsrc = zo.read('xl/workbook.xml').decode(), zs.read('xl/workbook.xml').decode()
wo2 = re.sub(r'<sheet name="门店余额核对[^>]*/>', '', wo)
wo2 = re.sub(r'<definedName name="_xlnm\.Print_(Titles|Area)" localSheetId="27">门店余额核对![^<]*</definedName>', '', wo2)
chk('workbook.xml 只多了两张表的登记和打印区域', wo2, wsrc)
chk('新表在第 28 张（打印区域的 localSheetId=27 对得上）', re.findall(r'<sheet name="([^"]+)"', wo)[27], SH)
chk('选中的工作表只有一张（WPS 不会进「工作组」）', sum(zo.read(n).count(b'tabSelected="1"') for n in zo.namelist()
                                           if n.startswith('xl/worksheets/sheet')), 1, 0)
rid = re.search(r'<sheet name="门店余额核对"[^>]*r:id="(rId\d+)"', wo).group(1)
main_part = 'xl/' + re.search(r'Id="%s"[^>]*Target="([^"]+)"' % rid,
                             zo.read('xl/_rels/workbook.xml.rels').decode()).group(1)
mx = zo.read(main_part).decode()
chk('新表有工作表保护（不设密码）', bool(re.search(r'<sheetProtection [^>]*sheet="1"', mx)) and 'password' not in mx, True)
chk('辅助表是隐藏的', bool(re.search(r'<sheet name="门店余额核对辅助"[^>]*state="hidden"', wo)), True)

print(f'核对 {N_OK[0] + len(BAD)} 项，一致 {N_OK[0]} 项，不一致 {len(BAD)} 项')
for b in BAD[:40]:
    print('  ✗', b)
sys.exit(1 if BAD else 0)
