# -*- coding: utf-8 -*-
"""【应收应付汇总】【往来对账单】【内部往来】：
   应收＝销项发票 − 收入类收款；应付＝计应付的进项发票 − 成本费用类付款；其他往来＝收支项目「往来」（借款、押金…）；
   内部往来＝5 家公司之间的「内部划转」。"""
import datetime as dt
from openpyxl.formatting.rule import FormulaRule
from common import *
from s_reports import year_crit

AR_R0 = 6
AR_N = PT_R1 - PT_R0 + 1        # 跟往来单位一行对一行


def _op(col, n, cc):
    return f'SUMIFS({br(col, OP_R0, OP_R1)},{br(OP_PARTY, OP_R0, OP_R1)},{n},{br(OP_CO, OP_R0, OP_R1)},{cc})'


def build_ar(wb, ctx):
    ws = wb.create_sheet(SH_AR)
    widths(ws, {'A': 6, 'B': 30, 'C': 8, 'D': 13, 'E': 13, 'F': 13, 'G': 14, 'H': 13, 'I': 13, 'J': 13, 'K': 14,
                'L': 14, 'M': 14, 'N': 20})
    title(ws, '应 收 应 付 汇 总（每个往来单位 · 按发票和资金台帐）', 'N', C_AR,
          '💡 选公司、年份。应收：期初 ＋ 本年开的销项发票 − 本年收到的收入类款项；应付：期初 ＋ 本年计应付的进项发票 − 本年付给供应商的成本费用类款项'
          '（税款、工资这类没有应付的不算）；'
          '其他往来：收支项目选「往来」的（借款、还款、押金），正数＝对方欠我方（其他应收），负数＝我方欠对方（其他应付）。'
          '净额＝应收 − 应付 ＋ 其他往来。应收是负数＝已收款还没开票（或预收）。点表头筛选按钮能只看有余额的。')
    selector(ws, 'B3', '公司', 'C3', '全部', f'={AUX_CO_ALL}')
    selector(ws, 'E3', '年份', 'F3', 2026, YEARS)
    cc, y = co_crit('$C$3'), '$F$3'
    y0, y1 = f'DATE({y},1,1)', f'DATE({y}+1,1,1)'
    header(ws, 4, [('A', '序号'), ('B', '往来单位'), ('C', '类型'), ('D', '期初应收'), ('E', '本年开票'), ('F', '本年收款'),
                   ('G', '应收余额'), ('H', '期初应付'), ('I', '本年进项\n（计应付）'), ('J', '本年付款'), ('K', '应付余额'),
                   ('L', '其他往来余额\n（借款押金等）'), ('M', '净额\n（对方欠我方＋）'), ('N', '提示')], C_AR)
    put(ws, 'A5', '合计', F_TXTB, FILL_TOT, align=AC)
    ws.merge_cells('A5:C5')
    for c in 'DEFGHIJKLM':
        put(ws, f'{c}5', f'=ROUND(SUM({c}{AR_R0}:{c}{AR_R0 + AR_N}),2)', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'N5', None, fill_=FILL_TOT)

    def cols(n, pc_cash, pc_inv):
        """n：单位名单元格；返回 D..M 的公式"""
        jb = lambda: f'{jr(J_DATE)},"<"&{y0}'
        jy = f'{year_crit(y)}'
        rec = lambda per: (f'SUMIFS({jr(J_INV)},{pc_cash},{jr(J_CLS)},"{CLS_IN}",{per})'
                           f'-SUMIFS({jr(J_OUTV)},{pc_cash},{jr(J_CLS)},"{CLS_IN}",{per})')
        pay = lambda per: (f'SUMIFS({jr(J_OUTV)},{pc_cash},{jr(J_APF)},1,{per})'
                           f'-SUMIFS({jr(J_INV)},{pc_cash},{jr(J_APF)},1,{per})')
        wl = (f'SUMIFS({jr(J_OUTV)},{pc_cash},{jr(J_CLS)},"{CLS_WL}",{jr(J_DATE)},"<"&{y1})'
              f'-SUMIFS({jr(J_INV)},{pc_cash},{jr(J_CLS)},"{CLS_WL}",{jr(J_DATE)},"<"&{y1})')
        vb = f'{vr(V_DATE)},"<"&{y0}'
        vy = f'{vr(V_DATE)},">="&{y0},{vr(V_DATE)},"<"&{y1}'
        return {
            'D': f'{_op(OP_AR, n, cc) if n else "0"}+SUMIFS({vr(V_ARV)},{pc_inv},{vb})-({rec(jb())})',
            'E': f'SUMIFS({vr(V_ARV)},{pc_inv},{vy})',
            'F': rec(jy),
            'H': f'{_op(OP_AP, n, cc) if n else "0"}+SUMIFS({vr(V_APV)},{pc_inv},{vb})-({pay(jb())})',
            'I': f'SUMIFS({vr(V_APV)},{pc_inv},{vy})',
            'J': pay(jy),
            'L': f'{_op(OP_OTH, n, cc) if n else "0"}+{wl}',
        }

    for i in range(AR_N):
        r = AR_R0 + i
        p = PT_R0 + i
        n = f'$B{r}'
        put(ws, f'A{r}', f'=IF({n}="","",{i + 1})', F_AUTO, align=AC)
        put(ws, f'B{r}', f'=IF({SH_PARTY}!$B${p}="","",{SH_PARTY}!$B${p})', F_AUTOB, align=AL)
        put(ws, f'C{r}', f'=IF({n}="","",{SH_PARTY}!$C${p}&"")', F_AUTO, align=AC)
        ne = f'$P{r}'           # 隐藏：名字转义后当条件用
        ws[f'P{r}'] = f'={esc(n)}'
        ws[f'P{r}'].font = F_HELP
        f = cols(ne, f'{jr(J_PARTY)},{ne},{jr(J_CO)},{cc}', f'{vr(V_PARTY)},{ne},{vr(V_CO)},{cc}')
        for c, v in f.items():
            put(ws, f'{c}{r}', f'=IF({n}="","",ROUND({v},2))', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{r}', f'=IF({n}="","",ROUND(D{r}+E{r}-F{r},2))', F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'K{r}', f'=IF({n}="","",ROUND(H{r}+I{r}-J{r},2))', F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'M{r}', f'=IF({n}="","",ROUND(G{r}-K{r}+L{r},2))', F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'N{r}', f'=IF({n}="","",IF(G{r}<0,"已收款未开票/预收",IF(K{r}<0,"多付/预付或票没到","")))', F_NOTE, align=AL)
    # 最后一行：没登记的单位（资金台帐、发票里有，往来单位表里没有的）合在一起，保证总数对得上
    r = AR_R0 + AR_N
    put(ws, f'A{r}', '', F_AUTO, align=AC)
    put(ws, f'B{r}', '（没登记的单位合计）', F_RED, FILL_AUTO, align=AL)
    put(ws, f'C{r}', '', F_AUTO, FILL_AUTO)
    tot = cols(None, f'{jr(J_PARTY)},"?*",{jr(J_CO)},{cc}', f'{vr(V_PARTY)},"?*",{vr(V_CO)},{cc}')
    def op_all(col):
        # 期初往来全部 − 对方是自家公司的（那些在【内部往来】里算）
        return (f'(SUMIFS({br(col, OP_R0, OP_R1)},{br(OP_CO, OP_R0, OP_R1)},{cc})'
                f'-SUMPRODUCT(({CO_NAMES}<>"")*SUMIFS({br(col, OP_R0, OP_R1)},{br(OP_CO, OP_R0, OP_R1)},{cc},{br(OP_PARTY, OP_R0, OP_R1)},{CO_NAMES})))')
    tot['D'] = op_all(OP_AR) + '+' + tot['D'].split('+', 1)[1]
    tot['H'] = op_all(OP_AP) + '+' + tot['H'].split('+', 1)[1]
    tot['L'] = op_all(OP_OTH) + '+' + tot['L'].split('+', 1)[1]
    for c, v in tot.items():
        put(ws, f'{c}{r}', f'=ROUND({v}-SUM({c}{AR_R0}:{c}{r - 1}),2)', F_RED, FILL_AUTO, MONEY, AR)
    put(ws, f'G{r}', f'=ROUND(D{r}+E{r}-F{r},2)', F_RED, FILL_AUTO, MONEY, AR)
    put(ws, f'K{r}', f'=ROUND(H{r}+I{r}-J{r},2)', F_RED, FILL_AUTO, MONEY, AR)
    put(ws, f'M{r}', f'=ROUND(G{r}-K{r}+L{r},2)', F_RED, FILL_AUTO, MONEY, AR)
    put(ws, f'N{r}', '到【数据校验】看是哪些，登记后会分到上面各行', F_NOTE, FILL_AUTO, align=AL)
    ws.auto_filter.ref = f'A4:N{r}'
    hide(ws, 'P')
    ws.freeze_panes = f'C{AR_R0}'
    return ws


# ─────────────────────────────── 往来对账单 ───────────────────────────────
ST_R0, ST_N = 8, 500
N_INV, N_CASH = 200, 300


def build_stmt(wb, ctx):
    ws = wb.create_sheet(SH_STMT)
    widths(ws, {'A': 11, 'B': 9, 'C': 30, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 15, 'I': 12, 'J': 12, 'K': 4,
                'L': 9, 'M': 12, 'N': 4, 'R': 12, 'S': 14})
    title(ws, '往 来 对 账 单（选公司、往来单位、起止日期 · 发票和收付款按日期排 · 每行带余额）', 'M', C_AR,
          '💡 C3 选公司（全部＝几家合起来），F3 选往来单位（客户、供应商、借款人、自家另一家公司都行），起止日期不填＝不限。'
          '开票（销项）让对方欠我方增加，对方付款减少；进项发票（计应付的）是我方欠对方，我方付款冲掉。'
          '余额正数＝对方欠我方，负数＝我方欠对方（或对方预付）。跟【应收应付汇总】一个口径：只列销项/进项发票、收入类收款、'
          '付给供应商的成本费用款、往来（借款押金）、内部划转；发工资、交税这类不算往来的不列。')
    selector(ws, 'B3', '公司', 'C3', '全部', f'={AUX_CO_ALL}')
    selector(ws, 'E3', '往来单位', 'F3', '宁波市华方塑料机械制造有限公司', f'={SH_AUX}!$C$1:$C$620')
    ws.merge_cells('F3:H3')
    selector(ws, 'I3', '起始日', 'J3', None, fmt=DATE)
    selector(ws, 'L3', '截止日', 'M3', None, fmt=DATE)
    ws['R3'] = '=IF(N($J$3)>0,$J$3,1)'
    ws['R4'] = '=IF(N($M$3)>0,$M$3,2958465)'
    ws['Q3'], ws['Q4'] = '起(实际)', '止(实际)'
    for c in ('Q3', 'Q4', 'R3', 'R4'):
        ws[c].font = F_HELP
    cc = co_crit('$C$3')
    n = '$F$3'
    ws['R5'] = f'={esc(n)}'
    ws['Q5'] = '单位(转义)'
    ws['Q5'].font = ws['R5'].font = F_HELP
    ne = '$R$5'
    s = '$R$3'
    # 第 4、5 行：摘要
    heads = [('C', '期初余额（起始日前）'), ('D', '本期开票'), ('E', '本期进项'), ('F', '本期收款'), ('G', '本期付款'),
             ('H', '期末余额')]
    for c, t in heads:
        put(ws, f'{c}4', t, F_KPI_L, fill('FFD9E1F2'), align=ACW)
    ws.row_dimensions[4].height = 30
    e = '$R$4'
    pc = f'{jr(J_PARTY)},{ne},{jr(J_CO)},{cc},{jr(J_ARF)},1'
    opening = (f'SUMIFS({br(OP_AR, OP_R0, OP_R1)},{br(OP_PARTY, OP_R0, OP_R1)},{ne},{br(OP_CO, OP_R0, OP_R1)},{cc})'
               f'-SUMIFS({br(OP_AP, OP_R0, OP_R1)},{br(OP_PARTY, OP_R0, OP_R1)},{ne},{br(OP_CO, OP_R0, OP_R1)},{cc})'
               f'+SUMIFS({br(OP_OTH, OP_R0, OP_R1)},{br(OP_PARTY, OP_R0, OP_R1)},{ne},{br(OP_CO, OP_R0, OP_R1)},{cc})'
               f'+SUMIFS({vr(V_SD)},{vr(V_DATE)},"<"&{s})-SUMIFS({vr(V_SE)},{vr(V_DATE)},"<"&{s})'
               f'+SUMIFS({jr(J_OUTV)},{pc},{jr(J_DATE)},"<"&{s})-SUMIFS({jr(J_INV)},{pc},{jr(J_DATE)},"<"&{s})')
    vin = f'{vr(V_DATE)},">="&{s},{vr(V_DATE)},"<="&{e}'
    jin = f'{pc},{jr(J_DATE)},">="&{s},{jr(J_DATE)},"<="&{e}'
    put(ws, 'C5', f'=IF({n}="","",ROUND({opening},2))', F_KPI_V, fmt=MONEY, align=AC)
    R1 = ST_R0 + ST_N - 1
    # 本期数直接按条件求和（不靠下面列出来的行，超出一屏也不会少算）
    for c, v in (('D', f'SUMIFS({vr(V_SD)},{vin})'), ('E', f'SUMIFS({vr(V_SE)},{vin})'),
                 ('F', f'SUMIFS({jr(J_INV)},{jin})'), ('G', f'SUMIFS({jr(J_OUTV)},{jin})')):
        put(ws, f'{c}5', f'=IF({n}="","",ROUND({v},2))', F_KPI_V, fmt=MONEY, align=AC)
    put(ws, 'H5', f'=IF({n}="","",ROUND(C5+D5-E5-F5+G5,2))', F_KPI_V, fmt=MONEY, align=AC)
    put(ws, 'I4', '结论', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('I5:M5')
    put(ws, 'I5', f'=IF({n}="","",IF(H5>0,"对方欠我方 "&TEXT(H5,"#,##0.00"),IF(H5<0,"我方欠对方（或对方预付）"&TEXT(-H5,"#,##0.00"),'
                  f'"两清")))', F_RED, align=AL)
    header(ws, ST_R0 - 1, [('A', '日期'), ('B', '类型'), ('C', '单号 / 摘要'), ('D', '开票（销项）'), ('E', '进项发票'),
                           ('F', '收到对方的款'), ('G', '付给对方的款'), ('H', '余额\n（对方欠我方＋）'), ('I', '收支项目'),
                           ('J', '账户 / 公司'), ('L', '来源'), ('M', '源行')], C_AR)
    # 隐藏：S 列前 200 格＝发票排序键，后 300 格＝资金台帐排序键；显示第 k 行＝合起来第 k 小
    ws['S' + str(ST_R0 - 1)] = '排序键'
    for i in range(N_INV):
        ws[f'S{ST_R0 + i}'] = f'=IFERROR(SMALL({vr(V_SKEY)},{i + 1}),"")'
    for i in range(N_CASH):
        ws[f'S{ST_R0 + N_INV + i}'] = f'=IFERROR(SMALL({jr(J_SKEY)},{i + 1}),"")'
    for r in range(ST_R0, R1 + 1):
        ws[f'S{r}'].font = F_HELP
    for i in range(ST_N):
        r = ST_R0 + i
        ws[f'T{r}'] = f'=IFERROR(SMALL($S${ST_R0}:$S${R1},{i + 1}),"")'
        ws[f'U{r}'] = f'=IF(T{r}="","",T{r}-INT(T{r}/100000)*100000)'           # 低位：<50000 发票行号，≥50000 资金台帐行号+50000
        ws[f'V{r}'] = f'=IF(U{r}="","",IF(U{r}>=50000,"资金","发票"))'
        ws[f'W{r}'] = f'=IF(U{r}="","",IF(U{r}>=50000,U{r}-50000,U{r}))'
        for c in 'TUVW':
            ws[f'{c}{r}'].font = F_HELP
        isv = f'V{r}="发票"'
        vi = lambda col: f'INDEX({SH_INV}!${col}:${col},W{r})'
        ji = lambda col: f'INDEX({SH_CASH}!${col}:${col},W{r})'
        put(ws, f'A{r}', f'=IF(T{r}="","",INT(T{r}/100000))', F_AUTO, fmt=DATE, align=AC)
        put(ws, f'B{r}', f'=IF(T{r}="","",IF({isv},IF(N({vi(V_SD)})<>0,"开票","进项发票"),IF(N({ji(J_INV)})>0,"收款","付款")))',
            F_AUTO, align=AC)
        put(ws, f'C{r}', f'=IF(T{r}="","",IF({isv},"发票 "&{vi(V_ENO)}&{vi(V_NO)}&IF({vi(V_STAT)}&""="正常",""," （"&{vi(V_STAT)}&"）"),'
                         f'TRIM({ji(J_MEMO)}&" "&{ji(J_ONAME)})))', F_AUTO, align=AL)
        put(ws, f'D{r}', f'=IF(T{r}="","",IF({isv},N({vi(V_SD)}),0))', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'E{r}', f'=IF(T{r}="","",IF({isv},N({vi(V_SE)}),0))', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{r}', f'=IF(T{r}="","",IF(NOT({isv}),N({ji(J_INV)}),0))', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{r}', f'=IF(T{r}="","",IF(NOT({isv}),N({ji(J_OUTV)}),0))', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'H{r}', f'=IF(T{r}="","",ROUND($C$5+SUM($D${ST_R0}:D{r})-SUM($E${ST_R0}:E{r})-SUM($F${ST_R0}:F{r})'
                         f'+SUM($G${ST_R0}:G{r}),2))', F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'I{r}', f'=IF(T{r}="","",IF({isv},{vi(V_USE)}&"",{ji(J_ITEM)}&""))', F_AUTO, align=AC)
        put(ws, f'J{r}', f'=IF(T{r}="","",IF({isv},{vi(V_CO)}&"",{ji(J_ACC)}&""))', F_AUTO, align=AC)
        put(ws, f'L{r}', f'=IF(T{r}="","",IF({isv},"发票导入","资金台帐"))', F_NOTE, align=AC)
        put(ws, f'M{r}', f'=IF(T{r}="","",W{r})', F_NOTE, fmt='0', align=AC)
    put(ws, f'A{R1 + 1}', f'="共 "&(COUNT({vr(V_SKEY)})+COUNT({jr(J_SKEY)}))&" 笔"&IF(OR(COUNT({vr(V_SKEY)})>{N_INV},COUNT({jr(J_SKEY)})>{N_CASH}),'
                          f'"，⚠ 超出一屏（发票最多 {N_INV} 张、收付款最多 {N_CASH} 笔），下面没列全，请缩短起止日期；上面的本期、期末数是全的","")',
        F_RED, align=AL, border=False)
    ws.conditional_formatting.add(f'A4:M5', FormulaRule(formula=[f'OR(COUNT({vr(V_SKEY)})>{N_INV},COUNT({jr(J_SKEY)})>{N_CASH})'],
                                                        fill=fill('FFFFEB9C')))
    hide(ws, 'Q', 'R', 'S', 'T', 'U', 'V', 'W')
    ws.freeze_panes = f'A{ST_R0}'
    ws.print_title_rows = f'1:{ST_R0 - 1}'
    return ws


# ─────────────────────────────── 内部往来 ───────────────────────────────
def build_intra(wb, ctx):
    ws = wb.create_sheet(SH_INTRA)
    CC = [CL(2 + i) for i in range(8)]     # B..I
    widths(ws, {'A': 14, **{c: 13 for c in CC}, 'J': 14, 'K': 3, 'L': 30})
    title(ws, '内 部 往 来（几家公司之间划转的钱 · 谁欠谁）', 'J', C_AR,
          '💡 资金台帐里对方户名是自家另一家公司的，自动记成「内部划转」。第一张表：行＝本公司，列＝对方公司，'
          '数＝期初往来 ＋ 本公司开给对方的票 − 对方开给本公司的票 ＋ 本公司转出去的 − 对方转进来的（正数＝对方欠本公司）。第二张表检查两边记账是不是一致（A 欠 B 的应该等于 B 借给 A 的），'
          '只导了一边的流水会显示差额。第三张是自家公司之间开的发票。截止日不填＝到现在。')
    selector(ws, 'A3', '截止日', 'B3', None, fmt=DATE)
    ws['L3'] = '=IF(N($B$3)>0,$B$3+1,2958466)'
    ws['L3'].font = F_HELP
    e = '$L$3'

    def grid(top, label, cell_formula):
        section(ws, top, 'A', 'J', label, C_AR)
        put(ws, f'A{top + 1}', '本公司 ＼ 对方', F_HDR, fill(C_AR), align=ACW)
        for i, c in enumerate(CC):
            put(ws, f'{c}{top + 1}', f'=IF({SH_BASE}!$B${CO_R0 + i}="","",{SH_BASE}!$B${CO_R0 + i})', F_HDR, fill(C_AR), align=ACW)
        put(ws, f'J{top + 1}', '合计', F_HDR, fill(C_AR), align=AC)
        for j in range(8):
            r = top + 2 + j
            put(ws, f'A{r}', f'=IF({SH_BASE}!$B${CO_R0 + j}="","",{SH_BASE}!$B${CO_R0 + j})', F_TXTB, FILL_SUB, align=AC)
            for i, c in enumerate(CC):
                v = cell_formula(f'$A{r}', f'{c}${top + 1}', r, c, top)
                put(ws, f'{c}{r}', f'=IF(OR($A{r}="",{c}${top + 1}="",$A{r}={c}${top + 1}),"",{v})', F_AUTO, fmt=MONEY, align=AR)
            put(ws, f'J{r}', f'=IF($A{r}="","",ROUND(SUM(B{r}:I{r}),2))', F_AUTOB, fmt=MONEY, align=AR)
        return top + 10

    def bal(a, b, *_):
        pc = f'{jr(J_CO)},{a},{jr(J_PARTY)},{b},{jr(J_CLS)},"{CLS_INTRA}",{jr(J_DATE)},"<"&{e}'
        opc = f'{br(OP_CO, OP_R0, OP_R1)},{a},{br(OP_PARTY, OP_R0, OP_R1)},{b}'
        op = (f'SUMIFS({br(OP_AR, OP_R0, OP_R1)},{opc})-SUMIFS({br(OP_AP, OP_R0, OP_R1)},{opc})'
              f'+SUMIFS({br(OP_OTH, OP_R0, OP_R1)},{opc})')
        inv = (f'SUMIFS({vr(V_EFF)},{vr(V_DIR)},"内部",{vr(V_SCO)},{a},{vr(V_BCO)},{b},{vr(V_DATE)},"<"&{e})'
               f'-SUMIFS({vr(V_EFF)},{vr(V_DIR)},"内部",{vr(V_SCO)},{b},{vr(V_BCO)},{a},{vr(V_DATE)},"<"&{e})')
        return f'ROUND({op}+{inv}+SUMIFS({jr(J_OUTV)},{pc})-SUMIFS({jr(J_INV)},{pc}),2)'

    t1 = 5
    nxt = grid(t1, '① 内部往来余额（正数＝对方欠本公司）', bal)

    def chk(a, b, r, c, top):
        # 对称位置：①表里 (a,b) 与 (b,a) 相加应为 0
        i = CI(c) - 2
        j = r - (top + 2)
        mirror = f'{CL(2 + j)}{t1 + 2 + i}'
        me = f'{c}{t1 + 2 + j}'
        return f'IF(ABS(N({me})+N({mirror}))<0.01,"✓","差 "&TEXT(N({me})+N({mirror}),"#,##0.00"))'
    nxt2 = grid(nxt + 1, '② 两边对得上吗（本公司记的 ＋ 对方记的 应该＝0）', chk)
    for r in range(nxt + 3, nxt + 11):
        for c in CC:
            ws[f'{c}{r}'].number_format = 'General'
            ws[f'{c}{r}'].alignment = AC

    def inv(a, b, *_):
        return (f'ROUND(SUMIFS({vr(V_EFF)},{vr(V_DIR)},"内部",{vr(V_SCO)},{a},{vr(V_BCO)},{b},{vr(V_DATE)},"<"&{e}),2)')
    grid(nxt2 + 1, '③ 自家公司之间开的发票（行＝开票方，列＝收票方）', inv)
    put(ws, f'A{nxt2 + 12}', '账户互转（同一家公司两个账户倒钱）应该一进一出抵成 0：', F_NOTE, align=AL, border=False)
    r = nxt2 + 13
    for j in range(8):
        co = f'{SH_BASE}!$B${CO_R0 + j}'
        put(ws, f'A{r + j}', f'=IF({co}="","",{co})', F_TXTB, FILL_SUB, align=AC)
        put(ws, f'B{r + j}', f'=IF(A{r + j}="","",ROUND(SUMIFS({jr(J_NET)},{jr(J_CO)},A{r + j},{jr(J_CLS)},"{CLS_XFER}",{jr(J_DATE)},"<"&{e}),2))',
            F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'C{r + j}', f'=IF(A{r + j}="","",IF(ABS(N(B{r + j}))<0.01,"✓","✗ 只记了一边？"))', F_AUTO, align=AC)
    ctx['intra_chk'] = (nxt + 3, nxt + 10, r, r + 7)
    hide(ws, 'L')
    return ws
