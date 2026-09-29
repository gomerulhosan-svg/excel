# -*- coding: utf-8 -*-
"""【资金台帐】全自动：把【流水1～流水8】【手工记账】里认出来的每一笔按账户依次排好（不用再粘到这里），
   算即时余额、认往来单位 / 客户 / 供应商 / 收支项目 / 所属公司。这张表只看不改——要改的在流水表那一行的手工列改
   （每行最右边「去改 →」点一下就跳过去），这样再粘新流水也不会错位。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.protection import SheetProtection
from common import *

# 往来对账单的选择格（工作簿 2 里的对账单用；工作簿 1 本身不放对账单键）
ST_CO, ST_PARTY, ST_S, ST_E = f"{SH_STMT}!$C$3", f"{SH_STMT}!$F$3", f"{SH_STMT}!$R$3", f"{SH_STMT}!$R$4"

MAP = f"{SH_AUX}!${AUX_MAP0}${AUX_R0}:$AB${AUX_R1}"
SRCS = ','.join(f'流水区{j}' for j in range(1, N_SRC + 1))
CUMS = ','.join(f'累计区{j}' for j in range(1, N_SRC + 1))


def stmt_key(ws, r):
    """往来对账单排序键（只在工作簿 2 的取数表 / 合并算数时放）"""
    g = lambda c: f'{c}{r}'
    per = (f'AND({g(J_PARTY)}={ST_PARTY},{g(J_PARTY)}<>"",OR({ST_CO}="全部",{g(J_CO)}={ST_CO}),'
           f'{g(J_DATE)}>={ST_S},{g(J_DATE)}<={ST_E},N({g(J_INV)})+N({g(J_OUTV)})<>0,N({g(J_ARF)})=1)')
    ws[f'{J_SKEY}{r}'] = f'=IF({per},{g(J_DATE)}*100000+50000+ROW(),"")'
    # 年月（报表按月汇总用一个数比日期区间快）
    ws[f'{J_YM}{r}'] = f'=IF(N({g(J_DATE)})=0,"",YEAR({g(J_DATE)})*100+MONTH({g(J_DATE)}))'
    ws[f'{J_COI}{r}'] = f'=IF({g(J_CO)}="",0,IFERROR(MATCH({g(J_CO)},{CO_NAMES},0),0))'   # 公司序号
    ws[f'{J_SKEY}{r}'].font = ws[f'{J_YM}{r}'].font = ws[f'{J_COI}{r}'].font = F_HELP


def build_cash(wb, ctx, with_stmt=False):
    ws = wb.create_sheet(SH_CASH)
    widths(ws, {'A': 6, 'B': 12, 'C': 17, 'D': 12, 'E': 12, 'F': 13, 'G': 19, 'H': 28, 'I': 16, 'J': 16,
                'K': 16, 'L': 10, 'M': 12, 'N': 11, 'O': 10, 'P': 26, 'Q': 8, 'R': 16, 'S': 14, 'T': 10, 'U': 9,
                'V': 14, 'W': 9, 'X': 32, 'Y': 7})
    title(ws, '资 金 台 帐（全自动：各流水表认出来的每一笔 · 即时余额 · 往来单位和收支项目自动认）', J_LAST_VIS, C_CASH,
          '💡 不用在这张表粘贴或打字：银行、微信、支付宝的流水粘到各自的【流水】表，现金等记在【手工记账】，这里自动一笔一笔排好。'
          '右边自动认：往来单位按【往来单位】里的名称 / 别名 / 账号认，收支项目按「手工 → 自家公司 → 绑定 → 摘要关键词 → 客户首笔新增、以后续费」的顺序定。'
          'X 列有 ✗ 的点筛选挑出来，点最右边「去改 →」跳到流水表那一行，在手工列选往来单位 / 收支项目（社保挂靠、费用类就在那里手工筛选）。'
          '这张表加了保护（没有密码）防止误改，筛选照常能用。')
    # 第 3 行：汇总条
    put(ws, 'A3', '待处理：', F_KPI_L, fill('FFD9E1F2'), align=AR)
    put(ws, 'B3', f'=COUNTIF({jr(J_CHK)},"✗*")&" 笔"', F_KPI_V, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('C3:D3')
    put(ws, 'C3', '已认出：', F_KPI_L, fill('FFD9E1F2'), align=AR)
    put(ws, 'E3', f'={AUX_TOTAL}', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'F3', '最后日期：', F_KPI_L, fill('FFD9E1F2'), align=AR)
    put(ws, 'G3', f'=IF(COUNT({jr(J_DATE)})=0,"",MAX({jr(J_DATE)}))', F_KPI_V, fill('FFD9E1F2'), DATE, AC)
    put(ws, 'H3', f'="容量 "&TEXT({AUX_TOTAL}/{J_R1 - J_R0 + 1},"0%")&"（{J_CAP} 笔）"', F_NOTE, fill('FFD9E1F2'), align=AL)
    put(ws, 'I3', '全部账户余额：', F_KPI_L, fill('FFD9E1F2'), align=AR)
    put(ws, 'J3', f'=ROUND(SUM({AC_OPENS})+SUM({jr(J_NET)}),2)', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    bands = [('A', 'B', '账户', 'FF305496'), ('C', 'J', '← 从流水表取来的原样内容', 'FF5B9BD5'),
             ('K', 'M', '流水表里手工改的', 'FFBF8F00'), ('N', J_LAST_VIS, '自动（不用填）', 'FF7F7F7F')]
    for c1, c2, t, col in bands:
        section(ws, 4, c1, c2, t, col)
    heads = [(J_SEQ, '序号'), (J_ACC, '账户'), (J_TIME, '交易时间'), (J_IN, '收入金额'), (J_OUT, '支出金额'),
             (J_BANKBAL, '银行余额\n（流水带的）'), (J_OACCT, '对方账号'), (J_ONAME, '对方户名'), (J_OBANK, '对方开户行'),
             (J_MEMO, '摘要'), (J_MPARTY, '往来单位\n（手工改的）'), (J_MITEM, '收支项目\n（手工改的）'), (J_NOTE, '备注'),
             (J_DATE, '日期'), (J_CO, '所属公司'), (J_PARTY, '往来单位'), (J_PTYPE, '往来类型'), (J_CUS, '客户'),
             (J_SUP, '供应商'), (J_ITEM, '收支项目'), (J_CLS, '收支类别'), (J_BAL, '账户余额\n（即时）'),
             (J_BCHK, '银行余额\n核对'), (J_CHK, '校验'), (J_GO, '改')]
    for col, t in heads:
        c = {'A': 'FF305496', 'B': 'FF305496'}.get(col)
        if c is None:
            c = 'FF5B9BD5' if CI(col) <= CI('J') else ('FFBF8F00' if CI(col) <= CI('M') else 'FF7F7F7F')
        put(ws, f'{col}{J_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[J_HDR].height = 34
    hidden = [(J_SRC, '源表'), (J_SROW, '源行'), (J_INV, '收入数'), (J_OUTV, '支出数'), (J_NET, '净额'), (J_TS, '时刻'),
              (J_AUTO, '自动匹配'), (J_NESC, '户名转义'), (J_PESC, '单位转义'), (J_KW, '关键词'), (J_APF, '冲应付'),
              (J_ARF, '算往来'), (J_TODO, '待登记名'), (J_TODO1, '首次'), (J_TODOC, '计数'), (J_DUPK, '查重键'),
              (J_BS, '块首行'), (J_BE, '块尾行'), (J_RIN, '原收入'), (J_ROUT, '原支出'), (J_RONE, '原金额'),
              (J_RFLAG, '原收支'), (J_RPAY, '原支付方式'), (J_RSTAT, '原状态'), (J_RNO, '原单号'), (J_RTIME, '原时间'),
              (J_SKIP, '跳过原因'), (J_FDIR, '收支方向'), (J_CIN, '算出收入'), (J_COUT, '算出支出'), (J_RBAL, '原余额'),
              (J_BASE, '余额基数')] \
        + ([(J_SKEY, '对账单键'), (J_YM, '年月'), (J_COI, '公司序号')] if with_stmt else [])
    for col, t in hidden:
        ws[f'{col}{J_HDR}'] = t
        ws[f'{col}{J_HDR}'].font = F_HELP

    PN, PTY, PB, PA1, PA2, PAC, POLD, PAP = (pr(PT_NAME), pr(PT_TYPE), pr(PT_BIND), pr(PT_N1), pr(PT_N2), pr(PT_ACCT),
                                             pr(PT_OLD), pr(PT_AP))
    PN0 = pr(PT_N0)
    KW, KWD, KWI = br(KW_WORD, KW_R0, KW_R1), br(KW_DIR, KW_R0, KW_R1), br(KW_ITEM, KW_R0, KW_R1)
    R0 = J_R0
    OFFS, CNTS = aux_rng(AUX_OFF), aux_rng(AUX_CNT)
    for r in range(J_R0, J_R1 + 1):
        g = lambda c: f'{c}{r}'
        z = g(J_SRC)
        k = f'(ROW()-{J_HDR})'
        mp = lambda fi: f'INDEX({MAP},{z},{fi})'
        rf = lambda n: f'IF({mp(FI[n])}=0,"",INDEX(CHOOSE({z},{SRCS}),{g(J_SROW)},{mp(FI[n])}))'
        f = {}
        f[J_SRC] = f'=IF({k}>{AUX_TOTAL},"",MATCH({k}-1,{OFFS},1))'
        f[J_SROW] = f'=IF({z}="","",MATCH({k}-1-INDEX({OFFS},{z}),CHOOSE({z},{CUMS}),1)+{S_HDR})'
        f[J_SEQ] = f'=IF({z}="","",{k})'
        f[J_ACC] = f'=IF({z}="","",IF({mp(FI["账户列"])}>0,{rf("账户列")}&"",INDEX({aux_rng(AUX_ACC)},{z})&""))'
        f[J_TIME] = f'=IF({z}="","",{rf("日期")})'
        f[J_RBAL] = f'=IF({z}="","",TRIM({rf("余额")}&""))'
        f[J_BANKBAL] = f'=IF({g(J_RBAL)}="","",{num(g(J_RBAL))})'
        f[J_OACCT] = f'=IF({z}="","",TRIM({rf("对方账号")}&""))'
        f[J_ONAME] = f'=IF({z}="","",TRIM({rf("对方户名")}&""))'
        f[J_OBANK] = f'=IF({z}="","",TRIM({rf("开户行")}&""))'
        # 摘要＋用途/备注；微信、支付宝没有的栏目导出成「/」，去掉
        f[J_MEMO] = f'=IF({z}="","",TRIM(SUBSTITUTE(" "&{rf("摘要")}&" "&{rf("用途/备注")}&" "," / "," ")))'
        f[J_MPARTY] = f'=IF({z}="","",TRIM({rf("往来单位(手工)")}&""))'
        f[J_MITEM] = f'=IF({z}="","",TRIM({rf("收支项目(手工)")}&""))'
        f[J_NOTE] = f'=IF({z}="","",{rf("备注(手工)")}&"")'
        f[J_RIN] = f'=IF({z}="","",{rf("收入")})'
        f[J_ROUT] = f'=IF({z}="","",{rf("支出")})'
        f[J_RONE] = f'=IF({z}="","",{rf("单列金额")})'
        f[J_RFLAG] = f'=IF({z}="","",TRIM({rf("收支标志")}&""))'
        f[J_RPAY] = f'=IF({z}="","",TRIM({rf("支付方式")}&""))'
        f[J_RSTAT] = f'=IF({z}="","",{rf("状态")}&"")'
        f[J_RNO] = f'=IF({z}="","",TRIM({rf("单号")}&""))'
        f[J_RTIME] = f'=IF({z}="","",{rf("时间")})'
        f[J_BS] = f'=IF({z}="","",INDEX({OFFS},{z})+{J_R0})'
        f[J_BE] = f'=IF({z}="","",INDEX({OFFS},{z})+INDEX({CNTS},{z})+{J_R0 - 1})'
        # 收入 / 支出：两列 / 一列＋收支标志 / 一列正收负支；微信支付宝「/」「不计收支」的，零钱提现记支出、充值记收入
        mode = f'INDEX({aux_rng(AUX_MODE)},{z})'
        fl = g(J_RFLAG)
        has = lambda *ws_: 'OR(' + ','.join(f'ISNUMBER(SEARCH("{w}",{fl}))' for w in ws_) + ')'
        # 收支标志：1＝收（收入/贷/入），-1＝支（支出/借/出），0＝都不是（「/」「不计收支」）
        f[J_FDIR] = f'=IF({fl}="",0,IF({has("收", "贷", "入")},IF({has("支", "借", "出")},0,1),IF({has("支", "借", "出")},-1,0)))'
        dr = g(J_FDIR)
        wd = f'ISNUMBER(SEARCH("提现",{g(J_MEMO)}))'
        cz = f'ISNUMBER(SEARCH("充值",{g(J_MEMO)}))'
        one = num(g(J_RONE))
        flip = f'INDEX({aux_rng(AUX_FLIP)},{z})=1'          # 借方＝进账的银行：两列对调
        rin, rout = f'IF({flip},{g(J_ROUT)},{g(J_RIN)})', f'IF({flip},{g(J_RIN)},{g(J_ROUT)})'
        f[J_CIN] = (f'=IF({z}="",0,ROUND(IF({mode}=1,ABS({num(rin)}),IF({mode}=2,IF({dr}=1,ABS({one}),'
                    f'IF(AND({dr}=0,{cz},NOT({wd})),ABS({one}),0)),IF({mode}=3,MAX({one},0),0))),2))')
        f[J_COUT] = (f'=IF({z}="",0,ROUND(IF({mode}=1,ABS({num(rout)}),IF({mode}=2,IF({dr}=-1,ABS({one}),'
                     f'IF(AND({dr}=0,{wd}),ABS({one}),0)),IF({mode}=3,MAX(-{one},0),0))),2))')
        f[J_IN] = f'=IF(OR({z}="",{g(J_SKIP)}<>"",{g(J_CIN)}=0),"",{g(J_CIN)})'
        f[J_OUT] = f'=IF(OR({z}="",{g(J_SKIP)}<>"",{g(J_COUT)}=0),"",{g(J_COUT)})'
        # 跳过：同一张流水表粘重了 / 交易关闭 / 用银行卡付的（银行流水里已经有）/ 不计收支
        f[J_DUPK] = (f'=IF({z}="","",{g(J_ACC)}&"|"&{g(J_TIME)}&"|"&{g(J_RIN)}&"|"&{g(J_ROUT)}&"|"&{g(J_RONE)}&"|"&{g(J_BANKBAL)}'
                     f'&"|"&LEFT({g(J_ONAME)},20)&"|"&{g(J_RNO)}&"|"&{g(J_RFLAG)}&"|"&LEFT({g(J_MEMO)},20))')
        dup = f'COUNTIF(INDEX(${J_DUPK}${J_R0}:${J_DUPK}${J_R1},{g(J_BS)}-{J_R0 - 1}):{g(J_DUPK)},{esc(g(J_DUPK))})>1'
        pay = g(J_RPAY)
        f[J_SKIP] = (f'=IF({z}="","",IF({dup},"同一张流水表里粘重了（只算一次）",'
                     f'IF(OR(ISNUMBER(SEARCH("关闭",{g(J_RSTAT)})),ISNUMBER(SEARCH("失败",{g(J_RSTAT)}))),"交易关闭/失败",'
                     f'IF(AND({pay}<>"",{pay}<>"/",ISERROR(SEARCH("零钱",{pay})),ISERROR(SEARCH("余额",{pay})),NOT({cz})),'
                     f'"用银行卡/信用卡付的（银行流水里已经有这一笔）",'
                     f'IF(AND({mode}=2,{dr}=0,NOT({wd}),NOT({cz})),"不计收支（零钱通、转给自己等）","")))))')
        f[J_INV] = f'=N({g(J_IN)})'
        f[J_OUTV] = f'=N({g(J_OUT)})'
        f[J_NET] = f'={g(J_INV)}-{g(J_OUTV)}'
        f[J_DATE] = f'=IF({z}="","",{date_parse(g(J_TIME))})'
        tfrac = lambda v: (f'IF(ISNUMBER({v}),MOD({v},1),IFERROR(TIMEVALUE(MID(TRIM({v}&""),FIND(" ",TRIM({v}&"")&" ")+1,8)),'
                           f'IFERROR(TIMEVALUE(TRIM({v}&"")),0)))')
        f[J_TS] = (f'=IF({g(J_DATE)}="","",ROUND(({g(J_DATE)}+IF({g(J_RTIME)}&""="",{tfrac(g(J_TIME))},{tfrac(g(J_RTIME))}))'
                   f'*86400,0))')
        f[J_CO] = f'=IF({g(J_ACC)}="","",IFERROR(INDEX({AC_COS},MATCH({g(J_ACC)},{AC_NAMES},0))&"",""))'
        # 往来单位：先认自家公司（全称），再认往来单位的名称 / 别名 1 / 别名 2，户名空着或都对不上再按对方账号认
        nh = g(J_NESC)
        f[J_NESC] = f'=IF({g(J_ONAME)}="","",{esc(norm(g(J_ONAME)))})'
        oacct = f'TRIM({g(J_OACCT)}&"")'
        byacct = f'IF({oacct}="","",IFERROR(INDEX({PN},MATCH({esc(oacct)},{PAC},0))&"",""))'
        f[J_AUTO] = (f'=IF(AND({nh}="",{oacct}=""),"",IF({nh}="",{byacct},'
                     f'IFERROR(INDEX({CO_NAMES},MATCH({nh},{CO_NORMS},0))&"",IFERROR(INDEX({PN},MATCH({nh},{PN0},0))&"",'
                     f'IFERROR(INDEX({PN},MATCH({nh},{PA1},0))&"",IFERROR(INDEX({PN},MATCH({nh},{PA2},0))&"",{byacct}))))))')
        f[J_PARTY] = f'=IF({g(J_MPARTY)}<>"",{g(J_MPARTY)},{g(J_AUTO)})'
        f[J_PESC] = f'={esc(g(J_PARTY))}'
        pe = g(J_PESC)
        f[J_PTYPE] = (f'=IF({g(J_PARTY)}="","",IF(ISNUMBER(MATCH({pe},{CO_NAMES},0)),"内部公司",'
                      f'IFERROR(INDEX({PTY},MATCH({pe},{PN},0))&"","未登记")))')
        f[J_CUS] = f'=IF({g(J_PTYPE)}="客户",{g(J_PARTY)},"")'
        f[J_SUP] = f'=IF({g(J_PTYPE)}="供应商",{g(J_PARTY)},"")'
        text = f'{g(J_ONAME)}&" "&{g(J_MEMO)}&" "&{g(J_NOTE)}'
        f[J_KW] = (f'=IF({g(J_INV)}+{g(J_OUTV)}=0,"",IFERROR(INDEX({KWI},MATCH(1,INDEX(({KW}<>"")*ISNUMBER(SEARCH({KW},{text}))'
                   f'*(({KWD}="全部")+({KWD}="")+({KWD}=IF({g(J_INV)}>0,"收入","支出"))>0),0),0)),""))')
        bind = f'IFERROR(INDEX({PB},MATCH({pe},{PN},0))&"","")'
        old = f'IFERROR(INDEX({POLD},MATCH({pe},{PN},0))&"","")'
        prior = (f'COUNTIFS({jr(J_CO)},{g(J_CO)},{jr(J_PARTY)},{pe},{jr(J_INV)},">0",{jr(J_DATE)},"<"&{g(J_DATE)})'
                 f'+COUNTIFS(${J_CO}${R0}:{g(J_CO)},{g(J_CO)},${J_PARTY}${R0}:{g(J_PARTY)},{pe},'
                 f'${J_INV}${R0}:{g(J_INV)},">0",${J_DATE}${R0}:{g(J_DATE)},{g(J_DATE)})-1')
        cusin = f'AND({g(J_PTYPE)}="客户",{g(J_INV)}>0)'      # 客户来款：新增/续费优先于摘要关键词
        f[J_ITEM] = (f'=IF({g(J_MITEM)}<>"",{g(J_MITEM)},IF(OR({g(J_ACC)}="",{g(J_INV)}+{g(J_OUTV)}=0),"",'
                     f'IF({g(J_PTYPE)}="内部公司",IF({g(J_PARTY)}={g(J_CO)},"账户互转","内部划转"),'
                     f'IF(AND({bind}<>"",OR({g(J_PTYPE)}<>"税务银行",{g(J_KW)}="")),{bind},IF(AND({g(J_KW)}<>"",NOT({cusin})),{g(J_KW)},'
                     f'IF(AND({cusin},{g(J_DATE)}<>""),'
                     f'IF(OR({old}="是",{prior}>0),"续费","新增"),""))))))')
        f[J_CLS] = f'=IF({g(J_ITEM)}="","",IFERROR(INDEX({IT_CLSS},MATCH({g(J_ITEM)},{IT_NAMES},0))&"","？"))'
        # 冲应付：成本费用 / 投资类的付款，且对方是「计应付」的单位（供应商、J 列填了是、或者有计应付的进项发票）
        apf = f'IFERROR(INDEX({PAP},MATCH({pe},{PN},0))&"","")'
        f[J_APF] = (f'=IF(AND(OR({g(J_CLS)}="{CLS_VAR}",{g(J_CLS)}="{CLS_FIX}",{g(J_CLS)}="{CLS_OTH}",{g(J_CLS)}="{CLS_INVT}"),'
                    f'{g(J_PARTY)}<>""),IF(OR(AND({g(J_PTYPE)}="供应商",{apf}<>"否"),{apf}="是",'
                    f'COUNTIFS({vr(V_PARTY)},{pe},{vr(V_USE)},"应付")>0),1,0),0)')
        f[J_ARF] = f'=IF(OR({g(J_CLS)}="{CLS_IN}",{g(J_CLS)}="{CLS_WL}",{g(J_CLS)}="{CLS_INTRA}",{g(J_APF)}=1),1,0)'
        # 即时余额：只在这个账户那一段里按时间先后算（倒序导出、后补的月份也对）。
        # 同一时刻有几笔（微信商户收款和手续费同一分钟、银行只有日期）：按导出顺序排（倒序导出就倒着排）；
        # 带银行余额的，这一笔的余额等于「之前 ＋ 这一笔」或「之前 ＋ 同一时刻全部」都算对上（同一分钟里谁先谁后导出文件不一定按顺序）
        seg = lambda c: f'INDEX(${c}${J_R0}:${c}${J_R1},{g(J_BS)}-{J_R0 - 1})'
        segE = lambda c: f'INDEX(${c}${J_R0}:${c}${J_R1},{g(J_BE)}-{J_R0 - 1})'
        seg2 = lambda c: f'{seg(c)}:{segE(c)}'
        f[J_BASE] = (f'=IF(OR({g(J_ACC)}="",{g(J_TS)}=""),"",IFERROR(INDEX({AC_OPENS},MATCH({g(J_ACC)},{AC_NAMES},0)),0)'
                     f'+SUMIFS({seg2(J_NET)},{seg2(J_ACC)},{g(J_ACC)},{seg2(J_TS)},"<"&{g(J_TS)}))')
        base, ts, acc, bank = g(J_BASE), g(J_TS), g(J_ACC), g(J_BANKBAL)
        desc = f'N({seg(J_TS)})>N({segE(J_TS)})'
        pre = f'SUMIFS({seg(J_NET)}:{g(J_NET)},{seg(J_ACC)}:{acc},{acc},{seg(J_TS)}:{ts},{ts})'
        suf = f'SUMIFS({g(J_NET)}:{segE(J_NET)},{acc}:{segE(J_ACC)},{acc},{ts}:{segE(J_TS)},{ts})'
        dft = f'IF({desc},{suf},{pre})'
        alltie = f'SUMIFS({seg2(J_NET)},{seg2(J_ACC)},{acc},{seg2(J_TS)},{ts})'
        f[J_BAL] = (f'=IF({base}="","",ROUND({base}+IF({bank}="",{dft},IF(COUNTIFS({seg2(J_TS)},{ts},{seg2(J_ACC)},{acc})<2,{dft},'
                    f'IF(ABS({base}+{g(J_NET)}-{bank})<0.005,{g(J_NET)},IF(ABS({base}+{alltie}-{bank})<0.005,{alltie},{dft})))),2))')
        f[J_BCHK] = (f'=IF(OR({g(J_BANKBAL)}="",{g(J_BAL)}="",{g(J_SKIP)}<>""),"",IF(ABS({g(J_BAL)}-{num(g(J_BANKBAL))})<0.005,"✓",'
                     f'"差 "&TEXT({g(J_BAL)}-{num(g(J_BANKBAL))},"#,##0.00")))')
        noparty_in = f'AND({g(J_PARTY)}="",{g(J_MITEM)}<>"",{g(J_CLS)}="{CLS_IN}")'
        f[J_TODO] = (f'=IF(OR({z}="",{g(J_SKIP)}<>""),"",IF(AND({g(J_PARTY)}="",{g(J_ONAME)}<>"",OR({g(J_ITEM)}="",{noparty_in})),{norm(g(J_ONAME))},'
                     f'IF({g(J_PTYPE)}="未登记",{g(J_PARTY)},"")))')
        f[J_TODO1] = f'=IF({g(J_TODO)}="",0,IF(MATCH({esc(g(J_TODO))},{jr(J_TODO)},0)=ROW()-{R0 - 1},1,0))'
        f[J_TODOC] = f'=N({J_TODOC}{r - 1})+{g(J_TODO1)}'
        f[J_CHK] = (f'=IF({z}="","",IF({g(J_SKIP)}<>"","跳过："&{g(J_SKIP)},'
                    f'IF({g(J_ACC)}="","✗ 流水表还没对应账户：到【基础资料】② 填账户名",'
                    f'IF(COUNTIF({AC_NAMES},{g(J_ACC)})=0,"✗ 账户不在基础资料里",'
                    f'IF({g(J_DATE)}="","✗ 日期看不懂（要像 2026/9/1 这样）",'
                    f'IF(AND({g(J_INV)}<>0,{g(J_OUTV)}<>0),"✗ 收入、支出只能有一边",'
                    f'IF({g(J_INV)}+{g(J_OUTV)}=0,"✗ 没有金额（或金额不是数字）",'
                    f'IF(AND({g(J_PARTY)}<>"",{g(J_PTYPE)}=""),"✗ 往来单位没选类型：到【往来单位】C 列选",'
                    f'IF({g(J_ITEM)}="",IF(AND({g(J_PARTY)}="",{g(J_ONAME)}<>""),"✗ 认不出：对方户名没登记，到【往来单位】加，或点「去改」手工选",'
                    f'"✗ 请点「去改」选收支项目"),'
                    f'IF({g(J_CLS)}="？","✗ 收支项目不在基础资料清单里",'
                    f'IF(LEFT({g(J_BCHK)},1)="差","⚠ 余额跟银行对不上（漏粘了流水，或期初余额不对？）",'
                    f'IF({g(J_PTYPE)}="未登记","⚠ 往来单位没在【往来单位】登记（统计照算）",'
                    f'IF({noparty_in},"⚠ 收入没选往来单位（客户统计、应收里算不到），点「去改」选一下",'
                    f'IF(AND({g(J_PTYPE)}="税务银行",{g(J_MITEM)}="",{g(J_ITEM)}={SH_BASE}!$U${IT_ROW["税金"]},{g(J_OUTV)}>0),'
                    f'"⚠ 税务扣款默认算税金：是社保（挂靠）的点「去改」改",'
                    f'IF({g(J_ITEM)}="新增","√ 新客户首笔","√")))))))))))))))')
        f[J_GO] = (f'=IF({z}="","",HYPERLINK("#\'"&INDEX({aux_rng(AUX_SHEET)},{z})&"\'!"&ADDRESS({g(J_SROW)},'
                   f'MAX(1,{mp(FI["往来单位(手工)"])}),4),"去改 →"))')
        for col, v in f.items():
            ws[f'{col}{r}'] = v
        if with_stmt:
            stmt_key(ws, r)
    # 样式
    src = [J_ACC, J_TIME, J_IN, J_OUT, J_BANKBAL, J_OACCT, J_ONAME, J_OBANK, J_MEMO]
    man = [J_MPARTY, J_MITEM, J_NOTE]
    auto = [J_SEQ, J_DATE, J_CO, J_PARTY, J_PTYPE, J_CUS, J_SUP, J_ITEM, J_CLS, J_BAL, J_BCHK, J_CHK, J_GO]
    fmts = {J_IN: MONEY, J_OUT: MONEY, J_BANKBAL: MONEY, J_BAL: MONEY, J_DATE: DATE, J_TIME: DTIME, J_OACCT: '@'}
    aligns = {J_ONAME: AL, J_OBANK: AL, J_MEMO: AL, J_MPARTY: AL, J_NOTE: AL, J_PARTY: AL, J_CUS: AL, J_SUP: AL,
              J_CHK: AL, J_IN: AR, J_OUT: AR, J_BANKBAL: AR, J_BAL: AR}
    style_rows(ws, J_R0, J_R1, src + man + auto, auto=src + man + auto, fmts=fmts, aligns=aligns,
               fills={**{c: FILL_PASTE for c in src}, **{c: fill('FFFFF7E0') for c in man}}, bold=[J_BAL])
    for r in range(J_R0, J_R1 + 1):
        for c in src:
            ws[f'{c}{r}'].fill = FILL_PASTE
        for c in man:
            ws[f'{c}{r}'].fill = fill('FFFFF7E0')
        ws[f'{J_GO}{r}'].font = Font(name=YH, sz=9, color='FF0563C1', underline='single')
    hid = [c for c, _ in hidden]
    for col in hid:
        for r in range(J_R0, J_R1 + 1):
            ws[f'{col}{r}'].font = F_HELP
        ws.column_dimensions[col].hidden = True
    rng = f'{J_CHK}{J_R0}:{J_CHK}{J_R1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},1)="✗"'], fill=FILL_WARN,
                                                   font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},1)="√"'],
                                                   font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},2)="跳过"'],
                                                   font=Font(name=YH, sz=10, color='FF808080')))
    ws.conditional_formatting.add(f'{J_ITEM}{J_R0}:{J_ITEM}{J_R1}',
                                  FormulaRule(formula=[f'${J_ITEM}{J_R0}="新增"'], font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.auto_filter.ref = f'A{J_HDR}:{J_LAST_VIS}{J_R1}'
    ws.freeze_panes = f'C{J_R0}'
    ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                    selectLockedCells=False, selectUnlockedCells=False)
    return ws
