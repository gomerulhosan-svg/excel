# -*- coding: utf-8 -*-
"""把你发来的三个文件读进来：农行流水（→ 资金台帐演示数据）、销项/进项发票（→ 发票导入演示数据），
   并整理出往来单位初始名单。只读原件，不改原件。"""
import os
import xlrd
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
REF = os.path.join(HERE, '..', '参考')
F_BANK = os.path.join(REF, '农行流水_2896_202609.xls')
F_SALE = os.path.join(REF, '销项发票_中智优力_202607-08.xlsx')
F_BUY = os.path.join(REF, '进项发票_中智优力_202607-08.xlsx')

MY_CO = ('中智优力', '中智优力（宁波）财务咨询有限公司', '91330212MA2CKKG076')
MY_ACC = '农行2896'
# 另外 4 家（你 9/29 发来的）
COMPANIES = [MY_CO,
             ('中慧优力', '中慧优力（宁波）税务师事务所有限公司', '91330212MADBKKTP3U'),
             ('琴馨', '宁波琴馨信息咨询有限公司', '92330212MAC3G3Q29L'),
             ('优茗', '宁波市鄞州五乡优茗食品商行（个体工商户）', '92330212MACGWL0EX0'),
             ('飞鲸', '飞鲸（宁波）企业咨询有限公司', '91330212MA2KQU3M12')]

# 流水2～流水8 对应的账户（流水2 工行新开户没有样表）；样表原样放进对应的流水表当演示
# (账户名, 所属公司, 类型, 账号, 开户行, 期初余额, 期初日期, 样表文件, 备注)
ACCOUNTS = [
    ('工行（新开户）', '', '银行', '', '中国工商银行', 0, None, None,
     '→【流水2】新开户还没流水：选所属公司（账户名可以改成 工行+尾号）'),
    ('鄞州农商9654', '中慧优力', '银行', '81170101302269654', '宁波鄞州农村商业银行五乡支行', 43581.20, (2026, 6, 1),
     '流水样表_中慧优力_鄞州农商9654_202606.xls', '→【流水3】这家银行借方＝进账，已按余额自动认出来'),
    ('农行8240', '优茗', '银行', '39408001040018240', '中国农业银行宁波五乡支行', 10746.61, (2026, 7, 1),
     '流水样表_优茗_农行8240_202607-09.xls', '→【流水4】'),
    ('农行7705', '琴馨', '银行', '39408001040017705', '中国农业银行', 61371.12, (2026, 6, 1),
     '流水样表_琴馨_农行7705_202606.xls', '→【流水5】'),
    ('飞鲸对公户', '飞鲸', '银行', '', '', 30058.07, (2026, 5, 1),
     '流水样表_飞鲸_银行_202605.xls', '→【流水6】按文件名猜是飞鲸的，请核对公司、补账号开户行'),
    ('微信-汪优丽', '中智优力', '微信', '', '', 0, (2026, 7, 1),
     '流水样表_微信个人_汪优丽_202607.xlsx', '→【流水7】个人微信：请核对算哪家公司；微信账单不带余额，期初填当时零钱余额'),
    ('微信商户0744', '中智优力', '微信', '1707250744', '', 42501.40, (2026, 9, 1),
     '流水样表_微信商户_中智优力_202609.xlsx', '→【流水8】微信支付商户号 1707250744 的资金账单'),
]


def bank_rows():
    """农行导出：第 3 行起是明细，列顺序 交易时间|收入金额|支出金额|账户余额|对方账号|对方户名|对方开户行|摘要。
       原样保留（交易时间是文本，金额是数字），跟你从导出文件复制粘贴过来一样。"""
    sh = xlrd.open_workbook(F_BANK).sheet_by_index(0)
    head = [sh.cell_value(1, c) for c in range(sh.ncols)]
    acct = str(head[0]).split(':', 1)[-1].strip()
    rows = []
    for r in range(3, sh.nrows):
        v = [sh.cell_value(r, c) for c in range(8)]
        if not str(v[0]).startswith('20'):
            continue
        out = []
        for i, x in enumerate(v):
            if i in (1, 2, 3):
                out.append(float(x) if x not in ('', None) else None)
            else:
                out.append(str(x) if x not in ('', None) else None)
        rows.append(out)
    first_in = rows[0][1] or 0
    first_out = rows[0][2] or 0
    opening = round(rows[0][3] - first_in + first_out, 2)
    return acct, opening, rows


def bank_raw():
    """农行导出文件原样（标题、账号行、表头、明细、末尾几行都在），给【流水1】当演示：就像你把导出文件整份复制粘贴过来"""
    sh = xlrd.open_workbook(F_BANK).sheet_by_index(0)
    out = []
    for r in range(sh.nrows):
        row = []
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            row.append(None if v in ('', None) else v)
        out.append(row)
    while out and all(v is None for v in out[-1]):
        out.pop()
    return out


def sample_raw(fname):
    """流水样表原样（xls 用 xlrd、xlsx 用 openpyxl 读第一张有内容的表）：数字还是数字，文字还是文字"""
    path = os.path.join(REF, fname)
    out = []
    if fname.endswith('.xls'):
        bk = xlrd.open_workbook(path)
        sh = next(s for s in bk.sheets() if s.nrows)
        for r in range(sh.nrows):
            out.append([None if v in ('', None) else v for v in (sh.cell_value(r, c) for c in range(sh.ncols))])
    else:
        ws = next(w for w in openpyxl.load_workbook(path).worksheets if w.max_row > 1)
        for row in ws.iter_rows(values_only=True):
            out.append([None if v in ('', None) else v for v in row])
    while out and all(v is None for v in out[-1]):
        out.pop()
    return out


def sample_payers():
    """样表里给几家公司打钱的单位（名字像公司 / 个体户的）→ 先登记成客户"""
    names = []
    for a in ACCOUNTS:
        if not a[7] or a[2] != '银行':
            continue
        rows = sample_raw(a[7])
        hdr = next(i for i, r in enumerate(rows) if any(str(v).strip() in ('交易日期', '交易时间') for v in r if v))
        h = [str(v or '').strip() for v in rows[hdr]]
        name_col = next(i for i, t in enumerate(h) if t in ('对方户名', '对方名称'))
        inc = [i for i, t in enumerate(h) if t in ('收入金额', '借方发生额')]      # 鄞州农商：借方＝进账
        for r in rows[hdr + 1:]:
            nm = str(r[name_col] or '').strip() if name_col < len(r) else ''
            amt = r[inc[0]] if inc and inc[0] < len(r) else None
            try:
                amt = float(str(amt).replace(',', '')) if amt not in (None, '') else 0
            except ValueError:
                amt = 0
            if amt > 0 and nm and any(k in nm for k in ('公司', '个体工商户', '商行')) \
                    and not any(k in nm for k in ('银行', '中智优力', '中慧优力', '琴馨', '优茗', '飞鲸', '利息')):
                nm = nm.replace('(', '（').replace(')', '）')
                if nm not in names:
                    names.append(nm)
    return names


def invoice_rows(path):
    """「发票基础信息」那一页，19 列原样；最后的「合计行」不要"""
    ws = openpyxl.load_workbook(path)['发票基础信息']
    out = []
    for row in ws.iter_rows(min_row=2, max_col=19, values_only=True):
        if row[0] in (None, '') or str(row[0]).startswith('合计'):
            continue
        out.append(list(row))
    return out


def parties(bank, sales, buys):
    """初始往来单位：销项的购买方＋流水里给我们打钱的单位 → 客户；看得出用途的几个付款对象 → 绑好收支项目。
       进项发票上的都是吃饭、加油、过路费这类报销票，不登记（不登记就不算应付）。"""
    cus = {}
    for r in sales:
        cus.setdefault(r[7].replace('(', '（').replace(')', '）'), r[6])
    for r in bank:
        if r[1] and r[5] and r[5] != '郭伟健':
            cus.setdefault(r[5], None)
    out = []
    for name, tax in cus.items():
        out.append(dict(名称=name, 类型='客户', 税号=tax))
    out.append(dict(名称='中国电信股份有限公司宁波分公司', 类型='供应商', 绑定='办公费用',
                    别名1='中国电信股份有限公司宁波分公司（电信）', 备注='公共缴费：电话宽带'))
    out.append(dict(名称='国库（税款）', 类型='税务银行', 绑定='税金',
                    别名1='待报解预算收入－国库信息处理系统扣税', 备注='银行扣税'))
    for nm in sample_payers():
        if nm not in cus:
            out.append(dict(名称=nm, 类型='客户', 备注='流水样表里来款的单位'))
    return out


if __name__ == '__main__':
    a, o, b = bank_rows()
    print(a, o, len(b), b[0])
    s, p = invoice_rows(F_SALE), invoice_rows(F_BUY)
    print(len(s), len(p))
    for x in parties(b, s, p):
        print(x)
