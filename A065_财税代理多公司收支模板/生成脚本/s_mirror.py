# -*- coding: utf-8 -*-
"""工作簿 2 的隐藏取数表：【资金台帐】【发票导入】【基础资料】【往来单位】跟工作簿 1 同名、同位置，
   每格都是 =[1]表名!格子（跨工作簿链接），报表的公式跟放在一本里时一模一样。
   往来对账单要按它自己的选择格排序，那几列（对账单键等）就在这里算。"""
from common import *
from s_cash import stmt_key
from s_inv import stmt_inv

LINK = '[1]'


def _cell(ws, sheet, col, r, text):
    ws[f'{col}{r}'] = f'={LINK}{sheet}!{col}{r}' + ('&""' if text else '')


def _sheet(wb, name):
    ws = wb.create_sheet(name)
    ws['A1'] = f'（取数用，别改：数据来自【{WB1_FILE}】的【{name}】）'
    ws['A1'].font = F_NOTE
    ws.sheet_state = 'hidden'
    return ws


def build_mirrors(wb):
    ws = _sheet(wb, SH_CASH)
    for r in range(J_R0, J_R1 + 1):
        for col in J_MIRROR:
            _cell(ws, SH_CASH, col, r, col not in J_MIRROR_NUM)
        stmt_key(ws, r)
    ws = _sheet(wb, SH_INV)
    for r in range(V_R0, V_R1 + 1):
        for col in V_MIRROR:
            _cell(ws, SH_INV, col, r, col not in V_MIRROR_NUM)
        stmt_inv(ws, r)
    ws = _sheet(wb, SH_BASE)
    for r in range(CO_R0, CO_R1 + 1):
        _cell(ws, SH_BASE, CO_NAME, r, True)
    for r in range(AC_R0, AC_R1 + 1):
        for col in (AC_NAME, AC_CO, AC_TYPE, AC_NO):
            _cell(ws, SH_BASE, col, r, True)
        for col in (AC_OPEN, AC_ODATE, AC_NOW, AC_LAST):
            _cell(ws, SH_BASE, col, r, False)
    for r in range(IT_R0, IT_R1 + 1):
        for col in (IT_NAME, IT_CLS, IT_SHOW):
            _cell(ws, SH_BASE, col, r, True)
    for r in range(OP_R0, OP_R1 + 1):
        for col in (OP_CO, OP_PARTY):
            _cell(ws, SH_BASE, col, r, True)
        for col in (OP_AR, OP_AP, OP_OTH):
            _cell(ws, SH_BASE, col, r, False)
    ws = _sheet(wb, SH_PARTY)
    for r in range(PT_R0, PT_R1 + 1):
        for col in (PT_NAME, PT_TYPE):
            _cell(ws, SH_PARTY, col, r, True)


def add_external_link(wb, sheet_names):
    """工作簿 2 → 工作簿 1 的链接（同一个文件夹里的相对路径）"""
    from openpyxl.workbook.external_link.external import (ExternalLink, ExternalBook, ExternalSheetNames,
                                                           ExternalSheetDataSet, ExternalSheetData)
    from openpyxl.packaging.relationship import Relationship
    eb = ExternalBook(sheetNames=ExternalSheetNames(sheetName=list(sheet_names)),
                      sheetDataSet=ExternalSheetDataSet(sheetData=[ExternalSheetData(sheetId=i) for i in range(len(sheet_names))]),
                      id='rId1')
    el = ExternalLink(externalBook=eb)
    el.file_link = Relationship(type='externalLinkPath', Target=WB1_FILE, TargetMode='External', Id='rId1')
    wb._external_links.append(el)
