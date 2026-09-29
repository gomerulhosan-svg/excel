# -*- coding: utf-8 -*-
"""工作簿 1：【数据校验】把「会让报表算错」的情况一项项数出来，再列出流水、发票里还没登记的单位；【首页】各账户情况、步骤、导航。
   工作簿 2：【首页】数据取到哪天、选公司年份看关键数、对账提示、报表导航。"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from s_reports import year_crit
from s_ar import AR_R0, AR_N

TODO_N = 60      # 待登记清单各列 60 个
LINKF = Font(name=YH, sz=11, bold=True, color='FF0563C1', underline='single')


def build_check(wb, ctx):
    ws = wb.create_sheet(SH_CHK)
    widths(ws, {'A': 6, 'B': 46, 'C': 12, 'D': 8, 'E': 50, 'F': 22, 'G': 14})
    title(ws, '数 据 校 验（全自动 · 每项应为 √）', 'F', C_CHK,
          '💡 这里把「会让报表算错」的情况一项项数出来。✗ 要改；⚠ 是提醒，看一眼没问题就行。改完回来看，变成 √ 就好了。'
          '下面还列出流水、发票里出现了、但【往来单位】里还没登记的名字：复制名字，到【往来单位】B 列最下面「选择性粘贴 → 数值」、选类型'
          '（要绑收支项目的顺手选上），资金台帐和应收应付就会自动认出来。')
    header(ws, 4, [('A', '序号'), ('B', '检查项'), ('C', '数量'), ('D', '状态'), ('E', '怎么处理'), ('F', '去哪看')], C_CHK, height=28)
    PN = pr(PT_NAME)
    ITN = IT_NAMES
    kw_i, kw_w = br(KW_ITEM, KW_R0, KW_R1), br(KW_WORD, KW_R0, KW_R1)
    op_p = br(OP_PARTY, OP_R0, OP_R1)
    imp_acc = f'{SH_AUX}!${AUX_ACC}${AUX_R0}:${AUX_ACC}${AUX_R0 + N_IMP - 1}'
    man_acc = f"'{SH_MAN}'!$A${S_R0}:$A${S_R1}"
    checks = [
        ('资金台帐：认不出往来单位 / 收支项目的笔数', f'=COUNTIF({jr(J_CHK)},"✗ 认不出*")+COUNTIF({jr(J_CHK)},"✗ 请点*")', 'E',
         '资金台帐筛选「校验」列，点「去改」到流水表那一行选：常来往的单位到【往来单位】登记（能绑项目的绑上）；社保挂靠、费用这类手工选', SH_CASH),
        ('资金台帐：其他 ✗（账户没对上、日期看不懂、金额不对…）', f'=COUNTIF({jr(J_CHK)},"✗*")-C5', 'E', '筛选「校验」列，按提示改', SH_CASH),
        ('流水表：粘了流水却没认出格式（找不到日期或金额列）', f'=SUM({SH_AUX}!$AD${AUX_R0}:$AD${AUX_R1})', 'E',
         '到那张流水表看第 3 行：把导出文件的表头行也粘进来，或者在第 6 行填每样东西在第几列', SH_CASH),
        ('资金台帐：算出的余额跟银行给的余额对不上', f'=COUNTIF({jr(J_CHK)},"⚠ 余额*")', 'W',
         '一般是漏粘了几行，或期初余额不对；从第一个对不上的那笔往前查', SH_CASH),
        ('资金台帐：手工选的往来单位没在【往来单位】登记', f'=COUNTIF({jr(J_CHK)},"⚠ 往来单位没*")', 'W',
         '统计照算；想在应收应付、客户统计里单独一行，就去【往来单位】登记', SH_CASH),
        ('资金台帐：收入只选了项目、没选往来单位', f'=COUNTIF({jr(J_CHK)},"⚠ 收入没*")', 'W',
         '收支汇总照算，但客户收入统计、应收里算不到这笔；点「去改」把客户选上（名单见下面）', SH_CASH),
        ('资金台帐：税务扣款（默认算税金）要确认是不是社保', f'=COUNTIF({jr(J_CHK)},"⚠ 税务扣款*")', 'W',
         '国库扣款里可能有社保、社保挂靠：是的话点「去改」改成人工-社保 / 社保挂靠，改了提醒就消失', SH_CASH),
        ('资金台帐：自动跳过的笔数（银行卡付的、交易关闭、粘重了、不计收支）', f'=COUNTIF({jr(J_CHK)},"跳过*")', 'I',
         '这些不算钱（银行流水里已经有，或者本来就不是收支）；筛选「校验」列「跳过」看一眼就行', SH_CASH),
        ('资金台帐：流水日期比账户的期初日期还早',
         f'=SUMPRODUCT(({AC_NAMES}<>"")*ISNUMBER({AC_ODATES})*COUNTIFS({jr(J_ACC)},{AC_NAMES},{jr(J_DATE)},"<"&{AC_ODATES}))', 'E',
         '【基础资料】② 把期初日期改成这个账户第一笔流水那天（期初余额也要是那一笔之前的余额），不然报表的期初期末会少算', SH_BASE),
        ('手工记账：导入账户的钱又在手工记账里记了',
         f'=SUMPRODUCT(({imp_acc}<>"")*(COUNTIF({man_acc},{imp_acc})>0))', 'W',
         '银行、微信、支付宝的钱只在它自己的流水表里记；手工记账只记现金等没有导出的账户（不然余额分两段算）', SH_MAN),
        ('发票导入：认不出是哪家公司的发票', f'=COUNTIF({vr(V_CHK)},"✗ 销方*")', 'E',
         '到【基础资料】① 公司，把几家公司的税号都填上（发票靠税号认公司）', SH_INV),
        ('发票导入：开票日期看不懂', f'=COUNTIF({vr(V_CHK)},"✗ 开票日期*")', 'E',
         '开票日期要像 2026-09-01 或 2026/9/1；看不懂的票不进应收应付和发票汇总', SH_INV),
        ('发票导入：重复粘贴的发票（只算一次）', f'=COUNTIF({vr(V_CHK)},"⚠ 重复*")', 'W', '不影响金额；想干净可以把重复的行删掉', SH_INV),
        ('发票导入：开了销项票、客户还没登记', f'=COUNTIF({vr(V_CHK)},"⚠ 客户没*")', 'W',
         '应收照算（在应收应付汇总最后一行「没登记的单位合计」）；登记后单独一行。名单见下面右边', SH_INV),
        ('基础资料：有流水的资金账户没选所属公司', f'=SUMPRODUCT(({AC_NAMES}<>"")*({AC_COS}="")*(COUNTIF({jr(J_ACC)},{AC_NAMES})>0))', 'E',
         '【基础资料】② 选所属公司（还没流水的新账户可以先空着，有了流水再选）', SH_BASE),
        ('基础资料：资金账户的所属公司不在公司清单里',
         f'=SUMPRODUCT(({AC_COS}<>"")*(COUNTIF({CO_NAMES},{AC_COS})=0))', 'E', '公司简称要跟 ① 里写的一模一样', SH_BASE),
        ('基础资料：资金账户名称重复', f'=SUMPRODUCT(({AC_NAMES}<>"")*(COUNTIF({AC_NAMES},{AC_NAMES})>1))', 'E',
         '每个账户起一个不一样的简称（比如 农行2896、微信-中智）', SH_BASE),
        ('基础资料：公司没填税号（这家的发票认不出来）', f'=SUMPRODUCT(({CO_NAMES}<>"")*({CO_TAXES}=""))', 'W',
         '【基础资料】① 补税号', SH_BASE),
        ('基础资料：收支项目没选类别', f'=SUMPRODUCT(({ITN}<>"")*({IT_CLSS}=""))', 'E', '类别决定落到报表哪一块，一定要选', SH_BASE),
        ('基础资料：收支项目名称重复', f'=SUMPRODUCT(({ITN}<>"")*(COUNTIF({ITN},{ITN})>1))', 'E',
         '两个项目不能同名（两个「其他」已经叫其他收入、其他成本）', SH_BASE),
        ('基础资料：摘要关键词对应的项目不在收支项目里',
         f'=SUMPRODUCT(({kw_w}<>"")*(COUNTIF({ITN},{kw_i})=0))', 'E', '④ 的收支项目从下拉选', SH_BASE),
        ('往来单位：名称重复', f'=SUMPRODUCT(({PN}<>"")*(COUNTIF({PN},{esc(PN)})>1))', 'E',
         '同一个单位只登记一行；别的写法填到别名 1、别名 2', SH_PARTY),
        ('往来单位：没选类型', f'=SUMPRODUCT(({PN}<>"")*({pr(PT_TYPE)}=""))', 'E',
         '【往来单位】C 列选客户 / 供应商 / 个人……，不选的话客户来款认不出新增、续费', SH_PARTY),
        ('往来单位：绑定的收支项目不在清单里',
         f'=SUMPRODUCT(({pr(PT_BIND)}<>"")*(COUNTIF({ITN},{pr(PT_BIND)})=0))', 'E', '从下拉选', SH_PARTY),
        ('往来单位：登记了自家公司（会被当成客户/供应商）',
         f'=SUMPRODUCT(({PN}<>"")*((COUNTIF({CO_NAMES},{esc(PN)})+COUNTIF({CO_FULLS},{esc(PN)}))>0))', 'E',
         '自家公司只在【基础资料】① 登记，从【往来单位】删掉，内部划转才认得出', SH_PARTY),
        ('期初往来：填了金额，没选所属公司或往来单位',
         f'=SUMPRODUCT(((({br(OP_AR, OP_R0, OP_R1)}<>0)+({br(OP_AP, OP_R0, OP_R1)}<>0)+({br(OP_OTH, OP_R0, OP_R1)}<>0))>0)'
         f'*((({br(OP_CO, OP_R0, OP_R1)}="")+({op_p}=""))>0))', 'E', '这几行哪张报表都算不进去，补上公司和单位', SH_BASE),
        ('期初往来：单位没在【往来单位】登记',
         f'=SUMPRODUCT(({op_p}<>"")*(COUNTIF({PN},{op_p})=0)*(COUNTIF({CO_NAMES},{op_p})=0))', 'W',
         '期初照算进合计；登记后才会在应收应付里单独一行', SH_BASE),
        (f'资金台帐：已认出笔数（共 {J_CAP} 笔）', f'={AUX_TOTAL}', 'C', f'超过 {J_CAP - 300} 时提醒：一年一本，或者找我把容量加大', SH_CASH),
        (f'流水表：最满的一张用到第几行（每张 {S_CAP} 行）', f'=MAX({SH_AUX}!$AE${AUX_R0}:$AE${AUX_R1})', 'C',
         f'超过 {S_CAP - 100} 时提醒：在那张表倒数几行中间右键「插入」若干行即可（不要在最后一行下面加）', SH_CASH),
        (f'发票导入：已用行数（{V_CAP} 行）', f'=COUNTA({vr(V_ENO)})+COUNTA({vr(V_NO)})-SUMPRODUCT(({vr(V_ENO)}<>"")*({vr(V_NO)}<>""))',
         'C', f'超过 {V_CAP - 200} 行时提醒：在倒数几行中间右键「插入」若干行，再把上面一行整行复制粘贴下来', SH_INV),
    ]
    caps = {f'资金台帐：已认出笔数（共 {J_CAP} 笔）': J_CAP - 300, f'流水表：最满的一张用到第几行（每张 {S_CAP} 行）': S_CAP - 100,
            f'发票导入：已用行数（{V_CAP} 行）': V_CAP - 200}
    R0 = 5
    for k, (lab, f, lvl, how, where) in enumerate(checks):
        r = R0 + k
        put(ws, f'A{r}', k + 1, F_AUTO, align=AC)
        put(ws, f'B{r}', lab, F_TXT, align=ALW)
        put(ws, f'C{r}', f, F_AUTOB, fmt=INT if lvl not in 'CI' else '0', align=AC)
        if lvl == 'C':
            put(ws, f'D{r}', f'=IF(C{r}>{caps[lab]},"⚠","√")', F_AUTOB, align=AC)
        elif lvl == 'I':
            put(ws, f'D{r}', '—', F_AUTOB, align=AC)
        else:
            put(ws, f'D{r}', f'=IF(C{r}=0,"√","{"✗" if lvl == "E" else "⚠"}")', F_AUTOB, align=AC)
        put(ws, f'E{r}', how, F_NOTE, align=ALW)
        c = put(ws, f'F{r}', where, Font(name=YH, sz=10, color='FF0563C1', underline='single'), align=AC)
        link(c, where)
        ws.row_dimensions[r].height = 30
    last = R0 + len(checks) - 1
    rng = f'D{R0}:D{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{R0}="✗"'], fill=FILL_WARN, font=Font(name=YH, sz=11, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{R0}="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'D{R0}="√"'], fill=FILL_OK, font=Font(name=YH, sz=11, bold=True, color='FF00B050')))
    s = last + 1
    put(ws, f'A{s}', '', F_TXTB, FILL_TOT)
    put(ws, f'B{s}', '要改的（✗）项数', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'C{s}', f'=COUNTIF(D{R0}:D{last},"✗")', F_RED, FILL_TOT, fmt='0"项"', align=AC)
    put(ws, f'D{s}', f'=IF(C{s}=0,"√","✗")', F_RED, FILL_TOT, align=AC)
    put(ws, f'E{s}', f'="提醒（⚠）"&COUNTIF(D{R0}:D{last},"⚠")&" 项"', F_NOTE, FILL_TOT, align=AL)
    put(ws, f'F{s}', None, fill_=FILL_TOT)
    ctx['chk_err_cell'] = f'{SH_CHK}!$C${s}'
    wb[SH_AUX]['B1'] = f'={SH_CHK}!$C${s}'

    # 待登记清单
    t = s + 2
    section(ws, t, 'A', 'C', f'资金台帐里认不出 / 没登记的对方户名（前 {TODO_N} 个）', C_CHK)
    section(ws, t, 'E', 'G', f'销项发票里没登记的客户（前 {TODO_N} 个）', C_CHK)
    header(ws, t + 1, [('A', '序号'), ('B', '对方户名（复制 → 【往来单位】B 列「选择性粘贴 → 数值」）'), ('C', '笔数'),
                       ('E', '客户名称（→ 往来单位 B 列，粘贴为数值）'), ('F', '纳税人识别号（→ H 列）'), ('G', '应收金额')], C_CHK, height=30)
    jt, jc = jr(J_TODO), jr(J_TODOC)
    vt, vc = vr(V_TODO), vr(V_TODOC)
    for k in range(1, TODO_N + 1):
        r = t + 1 + k
        put(ws, f'A{r}', f'=IF(B{r}="","",{k})', F_AUTO, align=AC)
        put(ws, f'B{r}', f'=IFERROR(INDEX({jt},MATCH({k},{jc},0)),"")', F_AUTOB, align=AL)
        put(ws, f'C{r}', f'=IF(B{r}="","",COUNTIF({jt},{esc(f"B{r}")}))', F_AUTO, fmt=INT, align=AC)
        put(ws, f'F{r}', f'=IFERROR(INDEX({vr(V_BTAX)},MATCH({k},{vc},0))&"","")', F_AUTO, fmt='@', align=AC)
        put(ws, f'E{r}', f'=IFERROR(INDEX({vt},MATCH({k},{vc},0)),"")', F_AUTOB, align=AL)
        put(ws, f'G{r}', f'=IFERROR(ROUND(SUMIFS({vr(V_ARV)},{vt},{esc(f"INDEX({vt},MATCH({k},{vc},0))")}),2),"")', F_AUTO, fmt=MONEY, align=AR)
    r = t + 2 + TODO_N
    put(ws, f'A{r}', f'="流水里共 "&MAX({jc})&" 个、发票里共 "&MAX({vc})&" 个没登记的名字"', F_NOTE, align=AL, border=False)
    ws.freeze_panes = 'A5'
    return ws


def _banner(ws, text, tip, h2=40):
    widths(ws, {c: 16 for c in 'ABCDEFGH'})
    ws.sheet_properties.tabColor = C_HOME[2:]
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:H1')
    put(ws, 'A1', text, Font(name=YH, sz=18, bold=True, color='FFFFFFFF'), fill(C_HOME), align=AC, border=False)
    ws.row_dimensions[1].height = 42
    ws.merge_cells('A2:H2')
    put(ws, 'A2', tip, F_TIP, FILL_TIP, align=ALW, border=False)
    ws.row_dimensions[2].height = h2


def _steps(ws, r0, head, steps):
    put(ws, f'A{r0}', head, F_SEC, fill(C_HOME), align=AL)
    ws.merge_cells(f'A{r0}:H{r0}')
    for i, (a, b) in enumerate(steps):
        r = r0 + 1 + i
        put(ws, f'A{r}', a, F_TXTB, FILL_SUB, align=AC)
        ws.merge_cells(f'B{r}:H{r}')
        put(ws, f'B{r}', b, F_TXT, align=ALW)
        for c in 'CDEFGH':
            ws[f'{c}{r}'].border = BD
        ws.row_dimensions[r].height = 34
    return r0 + 1 + len(steps)


def _nav(ws, n0, nav):
    put(ws, f'A{n0}', '表格导航（点表名直接跳过去）', F_SEC, fill(C_HOME), align=AL)
    ws.merge_cells(f'A{n0}:H{n0}')
    for c, t in (('A', '表名'), ('B', '干什么的'), ('F', '你要做什么')):
        put(ws, f'{c}{n0 + 1}', t, F_HDR, fill('FF595959'), align=AC)
    ws.merge_cells(f'B{n0 + 1}:E{n0 + 1}')
    ws.merge_cells(f'F{n0 + 1}:H{n0 + 1}')
    for i, (sh, what, todo) in enumerate(nav):
        r = n0 + 2 + i
        c = put(ws, f'A{r}', sh, LINKF, align=AC)
        link(c, sh)
        ws.merge_cells(f'B{r}:E{r}')
        put(ws, f'B{r}', what, F_TXT, align=ALW)
        ws.merge_cells(f'F{r}:H{r}')
        star = todo.startswith('★')
        put(ws, f'F{r}', todo, F_IN if star else F_NOTE, FILL_IN if star else None, align=ALW)
        for cc_ in 'CDEGH':
            ws[f'{cc_}{r}'].border = BD
            if star and cc_ in 'GH':
                ws[f'{cc_}{r}'].fill = FILL_IN
        ws.row_dimensions[r].height = 30
    return n0 + 2 + len(nav)


def build_home1(wb, ctx):
    """工作簿 1 首页：各流水表情况、步骤、导航"""
    ws = wb.create_sheet(SH_HOME)
    _banner(ws, '财税代理 · 多公司收支管理 ① 流水发票导入（所有录入都在这本）',
            f'💡 这本管录入：每个账户的流水粘到各自的【流水】表（网银/微信/支付宝导出的文件整份原样粘，按表头自动认格式），'
            f'现金等记【手工记账】，发票粘【发票导入】，全部自动汇进【资金台帐】。报表在另一本【{WB2_FILE}】，'
            f'两本放在同一个文件夹，打开报表那本时选「更新链接 / 启用内容」就是最新的（或者先打开这本再打开那本）。'
            f'淡黄格子手填，淡蓝是粘贴区，灰格子是公式别动。')
    put(ws, 'A4', '待处理（✗）', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B4', f'=COUNTIF({jr(J_CHK)},"✗*")&" 笔"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    put(ws, 'C4', '资金台帐', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'D4', f'={AUX_TOTAL}&" 笔"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    put(ws, 'E4', '最后一笔', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'F4', f'={AUX_LAST}', F_KPI_V, fill('FFF2F2F2'), DATE, AC)
    put(ws, 'G4', '数据校验 ✗', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'H4', f'={ctx["chk_err_cell"]}&" 项"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    ws.row_dimensions[4].height = 28
    # 各流水表
    t = 6
    put(ws, f'A{t}', '各账户流水表（点表名跳过去）', F_SEC, fill(C_HOME), align=AL)
    ws.merge_cells(f'A{t}:H{t}')
    header(ws, t + 1, [('A', '表名'), ('B', '账户'), ('C', '所属公司'), ('D', '已认出（笔）'), ('E', '最后一笔'), ('F', '当前余额'),
                       ('G', '待处理 ✗'), ('H', '格式')], 'FF595959', height=26)
    for j, nm in enumerate(SRC_SHEETS, 1):
        r = t + 1 + j
        a = f'{SH_AUX}!${AUX_ACC}${AUX_R0 + j - 1}'
        c = put(ws, f'A{r}', nm, LINKF, align=AC)
        link(c, nm)
        if j <= N_IMP:
            b = AC_R0 + j - 1
            put(ws, f'B{r}', f'=IF({a}="","（没填）",{a})', F_AUTOB, align=AC)
            put(ws, f'C{r}', f'={SH_BASE}!$J${b}&""', F_AUTO, align=AC)
            put(ws, f'E{r}', f'={SH_BASE}!$Q${b}', F_AUTO, fmt=DATE, align=AC)
            put(ws, f'F{r}', f'={SH_BASE}!$P${b}', F_AUTO, fmt=MONEY, align=AR)
            put(ws, f'H{r}', f"=IF({SH_AUX}!$AD${AUX_R0 + j - 1}=1,\"✗ 没认出\",IF({SH_AUX}!${AUX_CNT}${AUX_R0 + j - 1}>0,\"✓\",\"（还没流水）\"))",
                F_AUTO, align=AC)
        else:
            put(ws, f'B{r}', '现金等（每行选账户）', F_AUTO, align=AC)
            for c_ in 'CEF':
                put(ws, f'{c_}{r}', None)
            put(ws, f'H{r}', '固定格式', F_AUTO, align=AC)
        put(ws, f'D{r}', f'={SH_AUX}!${AUX_CNT}${AUX_R0 + j - 1}', F_AUTOB, fmt=INT, align=AC)
        put(ws, f'G{r}', f'=COUNTIFS({jr(J_SRC)},{j},{jr(J_CHK)},"✗*")', F_RED, fmt=INT, align=AC)
    r = t + 2 + N_SRC
    r = _steps(ws, r + 1, '怎么用（第一次按 ①～④ 设好，以后每月做 ⑤⑥，⑦ 看报表）', [
        ('① 基础资料', '填几家公司（简称、全称、税号）、资金账户：第 1～8 个账户对应【流水1】～【流水8】，第 9 个起是现金等手工记账的；'
                     '选所属公司；期初余额＝第一笔流水之前的余额，期初日期＝第一笔流水那天。收支项目可以往下加。'),
        ('② 往来单位', '常来往的客户、供应商登记一下：流水上的户名写法不一样的填到别名；房东、刻章店、电信这类绑上收支项目，以后自动归类。老客户「建账前已合作」填 是。'),
        ('③ 摘要关键词', '【基础资料】④：摘要里出现「费用外收」就是银行手续费、「扣税」就是税金……已经填了常见的，可以再加。'),
        ('④ 期初往来', '建账前就有的应收、应付、借款余额，在【基础资料】⑤ 填一次。'),
        ('⑤ 每月导流水', '每个账户导出的流水文件整份复制，粘到它的【流水】表下面第一个空行（第一次连表头一起粘，自动认格式）；'
                        '现金、代付的记【手工记账】。【资金台帐】校验列有 ✗ 的点「去改」，到流水表那一行手工选往来单位 / 收支项目（社保挂靠、费用类在这里筛）。'),
        ('⑥ 每月导发票', '电子税务局 → 发票查询统计 → 导出（销项、进项都导），打开后「发票基础信息」页的明细整块粘到【发票导入】A 列。几家公司的都粘一起，按税号自动分。'),
        ('⑦ 看报表', f'存盘后打开【{WB2_FILE}】（放同一个文件夹），选「更新链接」：收支汇总表、简易现金流量表、费用统计、客户收入统计、应收应付汇总、往来对账单、内部往来、资金余额表都在那本。'),
    ])
    _nav(ws, r + 1, [
        (SH_BASE, '公司、资金账户、收支项目（收支类型）、摘要关键词、期初往来', '★ 先设好，以后加账户/项目在这'),
        (SH_PARTY, '客户、供应商、个人、税务银行…… 绑定收支项目、流水户名别名', '★ 新客户/新供应商在这登记'),
        (SRC_SHEETS[0], '每个账户一张：导出的流水整份原样粘，自动认格式；认不出的在这一行手工选（流水1～流水8）', '★ 每月粘流水'),
        (SH_MAN, '现金、个人代付等没有导出文件的账户，一笔一行', '★ 有现金收付时记'),
        (SH_CASH, '所有账户的每一笔（自动）：即时余额、往来单位、客户、供应商、收支项目、所属公司', '看；✗ 的点「去改」'),
        (SH_INV, '电子税务局导出的购销发票（几家公司的都粘一起）', '★ 每月粘发票'),
        (SH_CHK, '所有数据问题一表列出；没登记的户名、客户名单', '有 ✗ 就按提示改'),
    ])
    return ws


def build_home2(wb, ctx, name=SH_HOME):
    """工作簿 2 首页：数据取到哪天、关键数、对账提示、报表导航"""
    ws = wb.create_sheet(name)
    _banner(ws, '财税代理 · 多公司收支管理 ② 汇总报表（只看，不用录入）',
            f'💡 这本的数都取自【{WB1_FILE}】（两本放同一个文件夹）。打开时 Excel/WPS 问「是否更新链接 / 启用内容」选「更新」，'
            f'或者先打开 ① 那本再打开这本，数就是最新的。每张报表顶上的黄格子选公司（或全部）、年份、日期。'
            f'这本里隐藏的「资金台帐 / 发票导入 / 基础资料 / 往来单位」是取数用的，别改。')
    put(ws, 'A4', '数据到', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B4', f'={AUX_LAST}', F_KPI_V, fill('FFF2F2F2'), DATE, AC)
    put(ws, 'C4', '资金台帐', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'D4', f'={AUX_NROW}&" 笔"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    put(ws, 'E4', '待处理（✗）', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'F4', f'=COUNTIF({jr(J_CHK)},"✗*")&" 笔"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    put(ws, 'G4', '录入表校验 ✗', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'H4', f'={AUX_ERR}&" 项"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    ws.row_dimensions[4].height = 28
    selector(ws, 'A6', '公司', 'B6', '全部', f'={AUX_CO_ALL}')
    selector(ws, 'D6', '年份', 'E6', 2026, YEARS)
    ws.row_dimensions[6].height = 24
    cc, y = co_crit('$B$6'), '$E$6'
    ci = coi_crit('$B$6')
    yc = year_crit(y)

    def net(cls, sign=+1):
        a, b = (J_INV, J_OUTV) if sign > 0 else (J_OUTV, J_INV)
        return (f'SUMIFS({jr(a)},{jr(J_CLS)},"{cls}",{jr(J_COI)},{ci},{yc})'
                f'-SUMIFS({jr(b)},{jr(J_CLS)},"{cls}",{jr(J_COI)},{ci},{yc})')
    kpis = [
        ('本年收入', f'=ROUND({net(CLS_IN)},2)', MONEY),
        ('变动利润', f'=ROUND({net(CLS_IN)}-({net(CLS_VAR, -1)}),2)', MONEY),
        ('利润', f'=ROUND({net(CLS_IN)}-({net(CLS_VAR, -1)})-({net(CLS_FIX, -1)}),2)', MONEY),
        ('本年新增客户', f'=COUNTIFS({jr(J_ITEM)},"新增",{jr(J_INV)},">0",{jr(J_COI)},{ci},{yc})', '0"个"'),
        ('资金余额（现在）', f'=ROUND(SUMIFS({AC_OPENS},{AC_COS},{cc})+SUMIFS({jr(J_NET)},{jr(J_COI)},{ci}),2)', MONEY),
        ('客户欠款＊', f'=ROUND(SUMIF({SH_AR}!$G${AR_R0}:$G${AR_R0 + AR_N},">0"),2)', MONEY),
        ('欠供应商＊', f'=ROUND(SUMIF({SH_AR}!$K${AR_R0}:$K${AR_R0 + AR_N},">0"),2)', MONEY),
        ('内部往来对不上', f'=COUNTIF({SH_INTRA}!$B${ctx["intra_chk"][0]}:$I${ctx["intra_chk"][1]},"差*")/2'
                        f'+COUNTIF({SH_INTRA}!$C${ctx["intra_chk"][2]}:$C${ctx["intra_chk"][3]},"✗*")', '0"处"'),
    ]
    for i, (lab, f, fmt) in enumerate(kpis):
        col = CL(i + 1)
        put(ws, f'{col}8', lab, Font(name=YH, sz=10, bold=True, color='FF595959'), fill('FFF2F2F2'), align=AC)
        put(ws, f'{col}9', f, Font(name=YH, sz=14, bold=True, color='FFC00000' if i < 7 else 'FF1F4E79'),
            fill('FFF2F2F2'), fmt, AC)
    ws.row_dimensions[9].height = 32
    ws.merge_cells('A10:H10')
    put(ws, 'A10', '收入、利润、新增客户按上面选的公司和年份；资金余额是所选公司所有账户到现在；'
                   '＊客户欠款、欠供应商按【应收应付汇总】里选的公司、年份（年底余额，只加正数）；内部往来对不上的看【内部往来】②和最下面。',
        F_NOTE, align=AL, border=False)
    _nav(ws, 12, [
        (SH_SUM, '「收支类型」× 1～12 月：收入、变动成本、变动利润、固定成本、利润', '选公司、年份看'),
        (SH_CF, '业务口径的简易现金流量表（经营 / 投资 / 筹资往来），跟账户余额核对', '选公司、年份看'),
        (SH_BAL, '每个账户期初、收、支、期末，跟银行余额核对', '选公司、起止日期看'),
        (SH_EXP, '收入、各项费用按公司对比', '选年份、月份范围看'),
        (SH_CUS, '每月新增客户、新增/续费收入；每个客户首次来款、续费、应收', '选公司、年份看'),
        (SH_AR, '每个往来单位的应收、应付、其他往来（借款押金）', '选公司、年份看'),
        (SH_STMT, '给某个客户/供应商的对账单：发票、收付款逐笔带余额', '选公司、单位、起止日期'),
        (SH_INTRA, '几家公司之间划转的钱谁欠谁、两边对不对得上；自家之间开的票', '选截止日看'),
        (SH_INVS, '每月销项、进项发票金额，对照同月收付款', '选公司、年份看'),
    ])
    return ws
