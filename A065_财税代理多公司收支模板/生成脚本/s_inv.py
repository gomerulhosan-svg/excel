# -*- coding: utf-8 -*-
"""【发票导入】电子税务局导出的「发票基础信息」整块粘贴（销项、进项、5 家公司都粘在一起），
   自动认本方公司（按税号）、销项/进项、往来单位，算应收、应付。"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from s_cash import ST_CO, ST_PARTY, ST_S, ST_E



def stmt_inv(ws, r):
    """往来对账单用的三列（只在工作簿 2 的取数表 / 合并算数时放）：这张票在对账单里算开票多少、进项多少、排序键。
       销项/进项按单位＋公司；自家公司之间的票，对方是所选单位（另一家自家公司）时也列上"""
    g = lambda c: f'{c}{r}'
    coall = f'{ST_CO}="全部"'
    ws[f'{V_SD}{r}'] = (f'=IF({g(V_DIR)}="销项",IF(AND({g(V_PARTY)}={ST_PARTY},OR({coall},{g(V_CO)}={ST_CO})),N({g(V_ARV)}),0),'
                        f'IF({g(V_DIR)}="内部",IF(AND({g(V_DUP)}="",{g(V_BCO)}={ST_PARTY},{g(V_SCO)}<>{ST_PARTY},'
                        f'OR({coall},{g(V_SCO)}={ST_CO})),N({g(V_EFF)}),0),0))')
    ws[f'{V_SE}{r}'] = (f'=IF({g(V_DIR)}="进项",IF(AND({g(V_PARTY)}={ST_PARTY},OR({coall},{g(V_CO)}={ST_CO})),N({g(V_APV)}),0),'
                        f'IF({g(V_DIR)}="内部",IF(AND({g(V_DUP)}="",{g(V_SCO)}={ST_PARTY},{g(V_BCO)}<>{ST_PARTY},'
                        f'OR({coall},{g(V_BCO)}={ST_CO})),N({g(V_EFF)}),0),0))')
    ws[f'{V_SKEY}{r}'] = (f'=IF(AND({g(V_SD)}+{g(V_SE)}<>0,{g(V_DATE)}>={ST_S},{g(V_DATE)}<={ST_E}),{g(V_DATE)}*100000+ROW(),"")')
    for c in (V_SD, V_SE, V_SKEY):
        ws[f'{c}{r}'].font = F_HELP


def build_inv(wb, ctx, with_stmt=False):
    ws = wb.create_sheet(SH_INV)
    wd = {V_SEQ: 5, V_CODE: 8, V_NO: 8, V_ENO: 22, V_STAX: 19, V_SNAME: 26, V_BTAX: 19, V_BNAME: 26, V_TIME: 17,
          V_AMT: 11, V_TAXAMT: 9, V_TOTAL: 11, V_SRC: 8, V_KIND: 10, V_STAT: 9, V_POS: 6, V_RISK: 6, V_ISSUER: 7,
          V_REM: 14, V_DATE: 11, V_CO: 10, V_DIR: 7, V_PARTY: 26, V_VAL: 12, V_USE: 8, V_MAP: 8, V_DUP: 7,
          V_ARV: 12, V_APV: 12, V_CHK: 26}
    widths(ws, wd)
    title(ws, '发 票 导 入（电子税务局「发票基础信息」整块粘贴 → 自动认公司、销项/进项、往来单位 → 应收应付）', V_CHK, C_INV,
          '💡 电子税务局 → 发票查询统计 → 导出，打开导出的文件，切到「发票基础信息」那一页，选中第 2 行起的明细（不要表头、不要最后的合计行）'
          '复制，粘到 A 列最下面的空行。销项、进项、几家公司的都粘在这一张里，按税号自动认是哪家公司、是销项还是进项。'
          '销项＝应收（客户欠我们的）；进项：【往来单位】里登记过的供应商算应付，没登记的（吃饭、加油、过路费这类报销小票）算「报销票」不计应付；'
          '个别要改的在 Z 列选「是 / 否」。红冲的蓝字票和红字票都留着，正负抵掉；作废的不算。同一张票粘了两次会标「重复」，只算一次。')
    section(ws, 3, V_SEQ, V_REM, '← 粘贴区：跟导出的「发票基础信息」一模一样的 19 列', 'FF5B9BD5')
    section(ws, 3, V_DATE, V_CHK, '自动（Z 列可以手工改）', C_INV)
    for i, t in enumerate(V_RAW):
        put(ws, f'{CL(i + 1)}{V_HDR}', t, F_HDR, fill('FF5B9BD5'), align=ACW)
    for col, t in ((V_DATE, '开票日期'), (V_CO, '本方公司'), (V_DIR, '方向'), (V_PARTY, '往来单位'), (V_VAL, '价税合计'),
                   (V_USE, '计入'), (V_MAP, '计应付\n（手工改）'), (V_DUP, '重复'), (V_ARV, '应收金额'), (V_APV, '应付金额'),
                   (V_CHK, '校验')):
        put(ws, f'{col}{V_HDR}', t, F_HDR, fill(C_INV if col != V_MAP else 'FFBF8F00'), align=ACW)
    ws.row_dimensions[V_HDR].height = 34
    for col, t in ((V_TODO, '待登记'), (V_TODO1, '首次'), (V_TODOC, '计数'), (V_SCO, '销方自家'),
                   (V_BCO, '购方自家'), (V_EFF, '有效价税'), (V_EAMT, '有效金额'), (V_ETAX, '有效税额'),
                   ) + (((V_SD, '对账开票'), (V_SE, '对账进项'), (V_SKEY, '对账单键')) if with_stmt else ()):
        ws[f'{col}{V_HDR}'] = t
        ws[f'{col}{V_HDR}'].font = F_HELP
    PN, PTAX, P0, P1, P2, PAP = pr(PT_NAME), pr(PT_TAX), pr(PT_N0), pr(PT_N1), pr(PT_N2), pr(PT_AP)
    R0 = V_R0
    for r in range(V_R0, V_R1 + 1):
        g = lambda c: f'{c}{r}'
        real = f'AND({g(V_ENO)}&{g(V_NO)}<>"",{g(V_SEQ)}&""<>"序号",LEFT({g(V_SEQ)}&"",2)<>"合计")'
        f = {}
        f[V_DATE] = f'=IF(NOT({real}),"",{date_parse(g(V_TIME))})'
        co = lambda tax, name: (f'IF({tax}&""="","",IFERROR(INDEX({CO_NAMES},MATCH({tax}&"",{CO_TAXES},0))&"",'
                                f'IFERROR(INDEX({CO_NAMES},MATCH({esc(norm(name))},{CO_NORMS},0))&"","")))')
        f[V_SCO] = f'=IF(NOT({real}),"",{co(g(V_STAX), g(V_SNAME))})'
        f[V_BCO] = f'=IF(NOT({real}),"",{co(g(V_BTAX), g(V_BNAME))})'
        f[V_CO] = f'=IF({g(V_SCO)}<>"",{g(V_SCO)},{g(V_BCO)})'
        f[V_DIR] = (f'=IF(NOT({real}),"",IF(AND({g(V_SCO)}<>"",{g(V_BCO)}<>""),"内部",IF({g(V_SCO)}<>"","销项",'
                    f'IF({g(V_BCO)}<>"","进项","未识别"))))')
        pty = lambda tax, name: (f'IFERROR(INDEX({PN},MATCH({tax}&"",{PTAX},0)),IFERROR(INDEX({PN},MATCH({esc(norm(name))},{P0},0)),'
                                 f'IFERROR(INDEX({PN},MATCH({esc(norm(name))},{P1},0)),IFERROR(INDEX({PN},MATCH({esc(norm(name))},{P2},0)),'
                                 f'{norm(name)}))))')
        f[V_PARTY] = (f'=IF({g(V_DIR)}="销项",{pty(g(V_BTAX), g(V_BNAME))},IF({g(V_DIR)}="进项",{pty(g(V_STAX), g(V_SNAME))},'
                      f'IF({g(V_DIR)}="内部",{g(V_BCO)},"")))')
        f[V_VAL] = f'=IF({g(V_DIR)}="","",IF(ISNUMBER(SEARCH("作废",{g(V_STAT)})),0,{num(g(V_TOTAL))}))'
        reg = f'ISNUMBER(MATCH({esc(g(V_PARTY))},{PN},0))'
        apflag = f'IFERROR(INDEX({PAP},MATCH({esc(g(V_PARTY))},{PN},0))&"","")'
        ptype = f'IFERROR(INDEX({pr(PT_TYPE)},MATCH({esc(g(V_PARTY))},{PN},0))&"","")'
        # 进项默认：登记成「供应商」的（J 列没填否）或 J 列填了「是」的算应付；其余（报销小票、登记成客户的饭店……）算报销票
        f[V_USE] = (f'=IF({g(V_DIR)}="销项","应收",IF({g(V_DIR)}="进项",IF({g(V_MAP)}="是","应付",IF({g(V_MAP)}="否","报销票",'
                    f'IF(OR(AND({ptype}="供应商",{apflag}<>"否"),{apflag}="是"),"应付","报销票"))),IF({g(V_DIR)}="内部","内部","")))')
        dupk = (f'IF({g(V_ENO)}<>"",COUNTIFS(${V_ENO}${R0}:{g(V_ENO)},{g(V_ENO)}&"*",${V_DIR}${R0}:{g(V_DIR)},{g(V_DIR)}),'
                f'COUNTIFS(${V_CODE}${R0}:{g(V_CODE)},{g(V_CODE)}&"*",${V_NO}${R0}:{g(V_NO)},{g(V_NO)}&"*",${V_DIR}${R0}:{g(V_DIR)},{g(V_DIR)}))')
        f[V_DUP] = f'=IF({g(V_DIR)}="","",IF({dupk}>1,"重复",""))'
        f[V_EFF] = f'=IF(AND({g(V_DIR)}<>"",{g(V_DIR)}<>"未识别",{g(V_DUP)}=""),{g(V_VAL)},0)'
        f[V_EAMT] = f'=IF({g(V_EFF)}=0,0,{num(g(V_AMT))})'
        f[V_ETAX] = f'=IF({g(V_EFF)}=0,0,{num(g(V_TAXAMT))})'
        f[V_ARV] = f'=IF(AND({g(V_USE)}="应收",{g(V_DUP)}=""),{g(V_VAL)},0)'
        f[V_APV] = f'=IF(AND({g(V_USE)}="应付",{g(V_DUP)}=""),{g(V_VAL)},0)'
        f[V_CHK] = (f'=IF({g(V_DIR)}="","",IF({g(V_DIR)}="未识别","✗ 销方、购方都不是自家公司：到【基础资料】核对公司税号",'
                    f'IF({g(V_DATE)}="","✗ 开票日期看不懂",IF({g(V_DUP)}<>"",IF({g(V_DIR)}="内部","内部开票：两家的导出里都有，只算一次","⚠ 重复粘贴了，只算一次"),'
                    f'IF(ISNUMBER(SEARCH("作废",{g(V_STAT)})),"作废票，不算",'
                    f'IF({g(V_DIR)}="内部","内部开票，看【内部往来】",IF({g(V_USE)}="报销票","报销票（不计应付）",'
                    f'IF(NOT({reg}),IF({g(V_DIR)}="销项","⚠ 客户没在【往来单位】登记（应收照算）","⚠ 供应商没在【往来单位】登记（应付照算）"),"√"))))))))')
        f[V_TODO] = f'=IF(AND({g(V_DIR)}="销项",NOT({reg})),{g(V_PARTY)},"")'
        f[V_TODO1] = f'=IF({g(V_TODO)}="",0,IF(MATCH({esc(g(V_TODO))},{vr(V_TODO)},0)=ROW()-{R0 - 1},1,0))'
        f[V_TODOC] = f'=N({V_TODOC}{r - 1})+{g(V_TODO1)}'
        for col, v in f.items():
            ws[f'{col}{r}'] = v
        if with_stmt:
            stmt_inv(ws, r)
    # 演示数据：你发来的销项 12 张、进项 97 张（原样）
    r = V_R0
    for row in ctx['sales'] + ctx['buys']:
        for i, v in enumerate(row):
            ws.cell(row=r, column=i + 1, value=v)
        r += 1
    raw = [CL(i) for i in range(1, 20)]
    auto = [V_DATE, V_CO, V_DIR, V_PARTY, V_VAL, V_USE, V_DUP, V_ARV, V_APV, V_CHK]
    style_rows(ws, V_R0, V_R1, raw + [V_MAP] + auto, auto=auto,
               fmts={V_AMT: MONEY, V_TAXAMT: MONEY, V_TOTAL: MONEY, V_VAL: MONEY, V_ARV: MONEY, V_APV: MONEY,
                     V_DATE: DATE, V_ENO: '@', V_STAX: '@', V_BTAX: '@'},
               aligns={V_SNAME: AL, V_BNAME: AL, V_PARTY: AL, V_CHK: AL, V_REM: AL},
               fills={**{c: FILL_PASTE for c in raw}, V_MAP: FILL_IN})
    for col in (V_TODO, V_TODO1, V_TODOC, V_SCO, V_BCO, V_EFF, V_EAMT, V_ETAX) + ((V_SD, V_SE, V_SKEY) if with_stmt else ()):
        for rr in range(V_R0, V_R1 + 1):
            ws[f'{col}{rr}'].font = F_HELP
        ws.column_dimensions[col].hidden = True
    for col in (V_CODE, V_NO, V_SRC, V_RISK, V_ISSUER):
        ws.column_dimensions[col].outlineLevel = 1
    rng = f'{V_CHK}{V_R0}:{V_CHK}{V_R1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${V_CHK}{V_R0},1)="✗"'], fill=FILL_WARN,
                                                   font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${V_CHK}{V_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    dv_list(ws, f'{V_MAP}{V_R0}:{V_MAP}{V_R1}', '"是,否"', '空着＝自动（登记过的供应商算应付）；是＝算应付；否＝报销票不算')
    ws.auto_filter.ref = f'A{V_HDR}:{V_CHK}{V_R1}'
    ws.freeze_panes = f'A{V_R0}'
    return ws
