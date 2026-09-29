# -*- coding: utf-8 -*-
"""A065 财税代理多公司收支模板 · 共用常量（表名 / 行列地址 / 容量）与样式。

各张表之间互相引用的行列地址都从这里取，不在各模块里手写，免得一处改了另一处没跟上。
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# ─────────────────────────── 表名 ───────────────────────────
SH_HOME = '首页'
SH_BASE = '基础资料'
SH_PARTY = '往来单位'
SH_CASH = '资金台帐'
SH_MAN = '手工记账'
SH_INV = '发票导入'
SH_SUM = '收支汇总表'
SH_CF = '简易现金流量表'
SH_BAL = '资金余额表'
SH_EXP = '费用统计'
SH_CUS = '客户收入统计'
SH_AR = '应收应付汇总'
SH_STMT = '往来对账单'
SH_INTRA = '内部往来'
SH_INVS = '发票汇总'
SH_CHK = '数据校验'
SH_AUX = '_辅助'

# 流水表：每个导入账户一张（整份粘网银 / 微信 / 支付宝导出的原样），最后一张手工记账（现金等）
N_IMP = 8                                           # 流水1～流水8 ↔ 基础资料 ② 第 1～8 个账户
SRC_SHEETS = [f'流水{i}' for i in range(1, N_IMP + 1)] + [SH_MAN]
N_SRC = len(SRC_SHEETS)

WB1_FILE = 'A065-1_流水发票导入.xlsx'
WB2_FILE = 'A065-2_汇总报表.xlsx'
WB1_ORDER = [SH_HOME, SH_BASE, SH_PARTY] + SRC_SHEETS + [SH_CASH, SH_INV, SH_CHK, SH_AUX]
WB2_REPORTS = [SH_SUM, SH_CF, SH_BAL, SH_EXP, SH_CUS, SH_AR, SH_STMT, SH_INTRA, SH_INVS]
SH_HOME2 = '报表首页'                               # 合并算数时用的名字；工作簿 2 里就叫「首页」
WB2_MIRRORS = [SH_CASH, SH_INV, SH_BASE, SH_PARTY]   # 工作簿 2 里隐藏的取数表，跟工作簿 1 同名同位置

# ─────────────────────────── 基础资料 ───────────────────────────
BASE_HDR = 4
CO_R0, CO_R1 = 5, 12            # 公司 8 行（现在 5 家）
CO_SEQ, CO_NAME, CO_FULL, CO_TAX, CO_NOTE, CO_NORM = 'A', 'B', 'C', 'D', 'E', 'F'   # F 隐藏：全称规范写法
AC_R0, AC_R1 = 5, 24            # 资金账户 20 行（你现在 15 个）
AC_SEQ, AC_NAME, AC_CO, AC_TYPE, AC_NO, AC_BANK, AC_OPEN, AC_ODATE, AC_NOW, AC_LAST, AC_NOTE = \
    'H', 'I', 'J', 'K', 'L', 'M', 'N', 'O', 'P', 'Q', 'R'
IT_R0, IT_R1 = 5, 64            # 收支项目 60 行
IT_SEQ, IT_NAME, IT_CLS, IT_SHOW, IT_NOTE = 'T', 'U', 'V', 'W', 'X'
KW_R0, KW_R1 = 5, 44            # 摘要关键词 40 行
KW_SEQ, KW_WORD, KW_DIR, KW_ITEM = 'Z', 'AA', 'AB', 'AC'
OP_R0, OP_R1 = 5, 204           # 期初往来 200 行
OP_SEQ, OP_CO, OP_PARTY, OP_AR, OP_AP, OP_OTH, OP_NOTE = 'AE', 'AF', 'AG', 'AH', 'AI', 'AJ', 'AK'

# 收支类别（跟你截图的分组一样；最后两个是系统用的）
CLS_IN, CLS_VAR, CLS_FIX, CLS_OTH, CLS_CAP, CLS_INVT, CLS_WL, CLS_INTRA, CLS_XFER = \
    '收入', '变动成本', '固定成本', '其他支出', '收到为准', '支出为准', '往来', '内部划转', '账户互转'
CLS_ALL = [CLS_IN, CLS_VAR, CLS_FIX, CLS_OTH, CLS_CAP, CLS_INVT, CLS_WL, CLS_INTRA, CLS_XFER]
CLS_COST = [CLS_VAR, CLS_FIX, CLS_OTH]          # 费用类（应付的付款也按这些算）

# 截图里的收支项目，一字不差；两个「其他」在下拉里要分得开，叫「其他收入」「其他成本」，报表上还显示「其他」
ITEMS = [
    (CLS_IN, ['项目', '续费', '新增', '代办', '刻章', '场地', '办证', '转介绍', '园区返税', '咨询服务', '其他收入']),
    (CLS_VAR, ['人工-工资', '人工-社保', '人工-福利', '税金', '差旅费', '培训费', '银行手续费', '刻章费用', '办公费用',
               '固定资产', '汽车费用', '交际支出', '佣金支出', '营销费用', '地址成本', '项目成本', '其他成本']),
    (CLS_FIX, ['店面租金', '装修成本']),
    (CLS_OTH, ['工资挂靠', '社保挂靠', '贷款利息']),
    (CLS_CAP, ['实收资本']),
    (CLS_INVT, ['投资']),
    (CLS_WL, ['往来']),
    (CLS_INTRA, ['内部划转']),
    (CLS_XFER, ['账户互转']),
]
SHOW_NAME = {'其他收入': '其他', '其他成本': '其他'}
# 每个收支项目在【基础资料】③ 的哪一行（报表按位置引用，基础资料里改名字报表跟着变）
IT_ROW = {}
for _cls, _names in ITEMS:
    for _n in _names:
        IT_ROW[_n] = 5 + len(IT_ROW)

# ─────────────────────────── 往来单位 ───────────────────────────
PT_HDR = 3
PT_R0, PT_R1 = 4, 603           # 600 个往来单位
(PT_SEQ, PT_NAME, PT_TYPE, PT_BIND, PT_AL1, PT_AL2, PT_ACCT, PT_TAX, PT_OLD, PT_AP, PT_NOTE,
 PT_N0, PT_N1, PT_N2) = 'A B C D E F G H I J K L M N'.split()      # L:N 隐藏：规范写法
PT_TYPES = ['客户', '供应商', '个人', '税务银行', '股东', '其他']

# ─────────────────────────── 流水表（每个账户一张） ───────────────────────────
S_LBL, S_AUTO, S_MAN, S_EFF = 4, 5, 6, 7            # 第 4～7 行：每样东西在第几列（标签 / 自动认的 / 手工改 / 实际用）
S_HDR = 8                                           # 第 8 行：列标题
S_R0, S_R1 = 9, 1008                                # 粘贴区 1000 行（一年的量；个人微信一个月六七十笔也够）
S_HSCAN = 40                                        # 表头只在粘贴区前 40 行里找
S_RAW = [CL(i) for i in range(1, 21)]               # A～T：原样粘贴
S_STAT, S_MP, S_MI, S_MN = 'U', 'V', 'W', 'X'       # 状态（自动）、往来单位 / 收支项目 / 备注（手工改）
S_HF, S_OK, S_CUM = 'Y', 'AA', 'AB'                 # 隐藏：是不是表头、是不是流水行、累计第几笔
S_X, S_B, S_VOTE = 'AC', 'AD', 'AE'                 # 隐藏：收入−支出、余额、这一行跟上一行的余额说明借贷方向正常(+1)/反了(-1)
S_NVOTE = 60                                        # 收支方向只看粘贴区前 60 行（一份导出文件的头几十笔就够判断）
S_DIRC = 'T'                                        # T4～T7：收支方向（自动 / 手工 / 实际）：正常 / 反向（借方＝进账的银行）
S_HROW = 'U'                                        # U4～U7：表头在第几行（自动 / 手工 / 实际）
# 字段（1～19），第 4～7 行 A～S 列一格一个
FIELDS = ['日期', '时间', '收入', '支出', '单列金额', '收支标志', '余额', '对方账号', '对方户名', '开户行',
          '摘要', '用途/备注', '支付方式', '状态', '单号', '往来单位(手工)', '收支项目(手工)', '备注(手工)', '账户列']
FI = {n: i + 1 for i, n in enumerate(FIELDS)}
N_DET = 15                                          # 前 15 个按表头自动认；后 4 个固定

# ─────────────────────────── 资金台帐（自动合并，只读） ───────────────────────────
J_HDR = 5
J_R0, J_R1 = 6, 3005            # 3000 笔（一年一本）
(J_SEQ, J_ACC, J_TIME, J_IN, J_OUT, J_BANKBAL, J_OACCT, J_ONAME, J_OBANK, J_MEMO,
 J_MPARTY, J_MITEM, J_NOTE,
 J_DATE, J_CO, J_PARTY, J_PTYPE, J_CUS, J_SUP, J_ITEM, J_CLS, J_BAL, J_BCHK, J_CHK, J_GO) = \
    'A B C D E F G H I J K L M N O P Q R S T U V W X Y'.split()
# 隐藏辅助列
(J_SRC, J_SROW, J_INV, J_OUTV, J_NET, J_TS, J_AUTO, J_NESC, J_PESC, J_KW, J_APF, J_ARF,
 J_TODO, J_TODO1, J_TODOC, J_DUPK, J_BS, J_BE, J_RIN, J_ROUT, J_RONE, J_RFLAG, J_RPAY, J_RSTAT, J_RNO, J_RTIME,
 J_SKIP, J_SKEY, J_FDIR, J_CIN, J_COUT, J_RBAL, J_BASE, J_YM, J_COI) = (
    'Z AA AB AC AD AE AF AG AH AI AJ AK AL AM AN AO AP AQ AR AS AT AU AV AW AX AY AZ BA BB BC BD BE BF BG BH').split()
J_CAP = J_R1 - J_R0 + 1
J_LAST_VIS = J_GO
J_MIRROR = [J_ACC, J_ONAME, J_MEMO, J_DATE, J_CO, J_PARTY, J_PTYPE, J_ITEM, J_CLS, J_BCHK, J_CHK,
            J_INV, J_OUTV, J_NET, J_APF, J_ARF]          # 工作簿 2 要取的列
J_MIRROR_NUM = {J_DATE, J_INV, J_OUTV, J_NET, J_APF, J_ARF}


def jr(col):
    return f"{SH_CASH}!${col}${J_R0}:${col}${J_R1}"


# ─────────────────────────── 发票导入 ───────────────────────────
V_HDR = 4
V_R0, V_R1 = 5, 2004            # 2000 张
V_CAP = V_R1 - V_R0 + 1
S_CAP = S_R1 - S_R0 + 1
# A..S：跟电子税务局导出的「发票基础信息」一模一样（整块粘贴）
V_RAW = ['序号', '发票代码', '发票号码', '数电发票号码', '销方识别号', '销方名称', '购方识别号', '购买方名称', '开票日期',
         '金额', '税额', '价税合计', '发票来源', '发票票种', '发票状态', '是否正数发票', '发票风险等级', '开票人', '备注']
(V_SEQ, V_CODE, V_NO, V_ENO, V_STAX, V_SNAME, V_BTAX, V_BNAME, V_TIME, V_AMT, V_TAXAMT, V_TOTAL,
 V_SRC, V_KIND, V_STAT, V_POS, V_RISK, V_ISSUER, V_REM) = [CL(i) for i in range(1, 20)]
(V_DATE, V_CO, V_DIR, V_PARTY, V_VAL, V_USE, V_MAP, V_DUP, V_ARV, V_APV, V_CHK) = \
    'T U V W X Y Z AA AB AC AD'.split()
V_SKEY = 'AE'                   # 隐藏：往来对账单排序键
V_TODO, V_TODO1, V_TODOC = 'AF', 'AG', 'AH'   # 隐藏：待登记单位
V_SCO, V_BCO = 'AI', 'AJ'       # 隐藏：销方/购方是不是自家公司
V_EFF, V_EAMT, V_ETAX = 'AK', 'AL', 'AM'   # 隐藏：有效（不重复、不作废）的价税合计 / 金额 / 税额
V_SD, V_SE = 'AN', 'AO'         # 隐藏：往来对账单里这张票算「开票」多少、算「进项」多少（按对账单选的公司、单位）
V_MIRROR = [V_NO, V_ENO, V_STAT, V_DATE, V_CO, V_DIR, V_PARTY, V_USE, V_DUP, V_ARV, V_APV, V_SCO, V_BCO,
            V_EFF, V_EAMT, V_ETAX]
V_MIRROR_NUM = {V_DATE, V_ARV, V_APV, V_EFF, V_EAMT, V_ETAX}


def vr(col):
    return f"{SH_INV}!${col}${V_R0}:${col}${V_R1}"


def br(col, r0, r1):
    return f"{SH_BASE}!${col}${r0}:${col}${r1}"


def pr(col):
    return f"{SH_PARTY}!${col}${PT_R0}:${col}${PT_R1}"


CO_NAMES = br(CO_NAME, CO_R0, CO_R1)
CO_FULLS = br(CO_FULL, CO_R0, CO_R1)
CO_NORMS = br(CO_NORM, CO_R0, CO_R1)
CO_TAXES = br(CO_TAX, CO_R0, CO_R1)
AC_NAMES = br(AC_NAME, AC_R0, AC_R1)
AC_COS = br(AC_CO, AC_R0, AC_R1)
AC_OPENS = br(AC_OPEN, AC_R0, AC_R1)
AC_ODATES = br(AC_ODATE, AC_R0, AC_R1)
IT_NAMES = br(IT_NAME, IT_R0, IT_R1)
IT_CLSS = br(IT_CLS, IT_R0, IT_R1)

# _辅助：下拉用的清单
AUX_CO_ALL = f"{SH_AUX}!$A$1:$A$9"       # 全部 + 8 个公司
# _辅助 E～AC：流水来源表（第 2～10 行一张表一行）：序号 表名 账户 笔数 偏移 字段1..19 取数方式
AUX_R0 = 2
AUX_R1 = AUX_R0 + N_SRC - 1
AUX_SEQ, AUX_SHEET, AUX_ACC, AUX_CNT, AUX_OFF = 'E', 'F', 'G', 'H', 'I'
AUX_MAP0 = 'J'                                   # J～AB：字段 1～19 的列号
AUX_MODE = 'AC'                                  # 1＝收入/支出两列 2＝一列金额＋收支标志 3＝一列金额正收负支
AUX_FLIP = 'AF'                                  # 1＝收入、支出两列反了（借方＝进账）
AUX_TOTAL = f"{SH_AUX}!${AUX_CNT}${AUX_R1 + 1}"
AUX_ERR = f"{SH_AUX}!$B$1"                       # 数据校验要改的项数（工作簿 2 首页用）
AUX_LAST = f"{SH_AUX}!$B$2"                      # 资金台帐最后一笔日期
AUX_NROW = f"{SH_AUX}!$B$3"                      # 资金台帐笔数


def aux_rng(col):
    return f"{SH_AUX}!${col}${AUX_R0}:${col}${AUX_R1}"


YEARS = '"2024,2025,2026,2027,2028,2029,2030"'


def norm(x):
    """名称规范写法：去空格、半角括号改全角（流水、发票里同一家公司括号写法常常不一样）"""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({x})," ",""),"(","（"),")","）")'


def co_crit(sel):
    """公司选择格 → SUMIFS 条件：选「全部」就是 "?*"（任何非空公司；"*" 在 Excel 里连公式算出的空文本也算进去）"""
    return f'IF({sel}="全部","?*",{sel})'


def coi_crit(sel):
    """公司选择格 → 资金台帐「公司序号」列的条件：全部＝">0"，单家＝它在 ① 里第几行（按数字比，比按名字 / 通配符快）"""
    return f'IF({sel}="全部",">0",IFERROR(MATCH({sel},{CO_NAMES},0),-1))'


def esc(x):
    """当 SUMIFS/COUNTIFS/MATCH 的条件用时，把名字里的 ~ * ? 转义成普通字符"""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'


# ─────────────────────────── 样式（跟 A064 一个样：微软雅黑，彩色标题带 ＋ 💡提示） ───────────────────────────
YH = '微软雅黑'
F_TITLE = Font(name=YH, sz=16, bold=True, color='FFFFFFFF')
F_TIP = Font(name=YH, sz=10, color='FFFF6600')
F_HDR = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_TXT = Font(name=YH, sz=10, color='FF000000')
F_TXTB = Font(name=YH, sz=10, bold=True, color='FF000000')
F_IN = Font(name=YH, sz=10, color='FF1F4E79')
F_AUTO = Font(name=YH, sz=10, color='FF404040')
F_AUTOB = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_NOTE = Font(name=YH, sz=9, color='FF808080')
F_HELP = Font(name=YH, sz=8, color='FFBFBFBF')
F_RED = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_KPI_L = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_KPI_V = Font(name=YH, sz=13, bold=True, color='FFC00000')
F_SEC = Font(name=YH, sz=11, bold=True, color='FFFFFFFF')
F_SEL = Font(name=YH, sz=11, bold=True, color='FF000000')


def fill(rgb):
    return PatternFill('solid', fgColor=rgb)


C_BASE, C_CASH, C_CONV, C_INV, C_RPT, C_AR, C_CHK, C_HOME = \
    'FF595959', 'FF2F75B5', 'FF7F7F7F', 'FF7030A0', 'FF375623', 'FFBF8F00', 'FF833C0C', 'FF1F3864'
C_IN, C_VAR, C_FIX, C_OTH, C_MISC = 'FFE2EFDA', 'FFDDEBF7', 'FFFFF2CC', 'FFF2F2F2', 'FFFFFFFF'   # 截图里各组的底色
C_IN_D, C_VAR_D, C_FIX_D = 'FF70AD47', 'FF5B9BD5', 'FFBF8F00'

FILL_TIP = fill('FFFFF2CC')
FILL_IN = fill('FFFFF7E0')      # 手工录入：淡黄
FILL_PASTE = fill('FFEAF1FB')   # 粘贴区：淡蓝
FILL_AUTO = fill('FFF2F2F2')    # 自动：淡灰
FILL_SUB = fill('FFDDEBF7')     # 小计
FILL_TOT = fill('FFFCE4D6')     # 合计 / 利润
FILL_SEL = fill('FFFFFF00')     # 选择格（亮黄）
FILL_OK = fill('FFE2EFDA')
FILL_WARN = fill('FFFFC7CE')
FILL_NONE = PatternFill(fill_type=None)

thin = Side(style='thin', color='FFBFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBD = Border()
AC = Alignment(horizontal='center', vertical='center')
ACW = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL = Alignment(horizontal='left', vertical='center')
ALW = Alignment(horizontal='left', vertical='center', wrap_text=True)
AR = Alignment(horizontal='right', vertical='center')

MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
MONEY0 = '#,##0;[Red]-#,##0;"-"'
DATE = 'yyyy/mm/dd'
DTIME = 'yyyy/mm/dd hh:mm'
INT = '0;-0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
MONTH = 'm"月"'


def put(ws, coord, value=None, font=None, fill_=None, fmt=None, align=None, border=True):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if fill_ is not None:
        c.fill = fill_
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border:
        c.border = BD
    return c


def title(ws, text, last_col, color, tip=None, h1=33, h2=None):
    """第 1 行标题带 ＋ 第 2 行 💡 提示；提示长就把第 2 行加高"""
    if h2 is None:
        width = sum((ws.column_dimensions[CL(i)].width or 9) for i in range(1, CI(last_col) + 1))
        per_line = max(20, int(width / 1.9))
        lines = -(-len(tip or '') // per_line)
        h2 = max(24, 16 * lines + 6)
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', text, F_TITLE, fill(color), align=AC, border=False)
    ws.row_dimensions[1].height = h1
    if tip:
        ws.merge_cells(f'A2:{last_col}2')
        put(ws, 'A2', tip, F_TIP, FILL_TIP, align=ALW, border=False)
        ws.row_dimensions[2].height = h2
    ws.sheet_properties.tabColor = color[2:] if len(color) == 8 else color
    ws.sheet_view.showGridLines = False


def header(ws, row, cols_texts, color, font=F_HDR, height=34):
    for col, t in cols_texts:
        put(ws, f'{col}{row}', t, font, fill(color), align=ACW)
    ws.row_dimensions[row].height = height


def section(ws, row, c1, c2, text, color):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    put(ws, f'{c1}{row}', text, F_SEC, fill(color), align=AL)
    for i in range(CI(c1) + 1, CI(c2) + 1):
        ws.cell(row=row, column=i).border = BD


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def hide(ws, *cols):
    for c in cols:
        ws.column_dimensions[c].hidden = True


def link(cell, sheet, ref='A1'):
    """表内跳转：写成 location，WPS / Excel 都能点过去"""
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{ref}",
                               display=None if cell.value is None else str(cell.value))
    return cell


def style_rows(ws, r0, r1, cols, auto=(), fmts=None, aligns=None, fills=None, bold=()):
    """数据区：白底细灰框；公式列淡灰；fills 可给手工列指定底色"""
    fmts, aligns, fills = fmts or {}, aligns or {}, fills or {}
    for r in range(r0, r1 + 1):
        for col in cols:
            c = ws[f'{col}{r}']
            is_auto = col in auto
            c.font = F_AUTOB if col in bold else (F_AUTO if is_auto else F_IN)
            c.fill = FILL_AUTO if is_auto else fills.get(col, FILL_NONE)
            c.border = BD
            c.alignment = aligns.get(col, AC)
            if col in fmts:
                c.number_format = fmts[col]


def dv_list(ws, sqref, formula, prompt=None, stop=True, blank=True):
    dv = DataValidation(type='list', formula1=formula, allow_blank=blank, showErrorMessage=stop,
                        errorStyle='stop' if stop else 'warning')
    if prompt:
        dv.promptTitle, dv.prompt = '提示', prompt
        dv.showInputMessage = True
    dv.errorTitle = '不在清单里'
    dv.error = '请从下拉里选；清单里没有的，先到【基础资料】/【往来单位】里登记'
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def dv_date(ws, sqref):
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050', allow_blank=True,
                        showErrorMessage=True, errorStyle='warning', errorTitle='日期', error='要像 2026/9/1 这样填日期')
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def selector(ws, cell_lbl, lbl, cell_in, value, dv_formula=None, fmt=None, prompt=None):
    """报表顶上的选择格：标签 ＋ 亮黄输入格"""
    put(ws, cell_lbl, lbl, F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, cell_in, value, F_SEL, FILL_SEL, fmt=fmt, align=AC)
    if dv_formula:
        dv = DataValidation(type='list', formula1=dv_formula, allow_blank=True, showErrorMessage=False)
        if prompt:
            dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', prompt, True
        ws.add_data_validation(dv)
        dv.add(cell_in.replace('$', ''))


def date_parse(x):
    """把粘贴进来的各种日期写法变成真日期（取日期部分）：
       真日期/日期时间 → 取整；"2026-09-01 08:24:04" / "2026/9/1" / "2026.9.1" / "2026年9月1日" / "20260901" 文本 → 拆年月日。
       只有纯 8 位数字才按 yyyymmdd 拆（"2026-9-1" 也是 8 个字，不能按位置拆）；小于 2000 年的数（比如只有时间）算看不懂"""
    s = (f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({x}),"/","-"),".","-"),"年","-"),"月","-"),"日"," ")')
    dp = f'LEFT({s},FIND(" ",{s}&" ")-1)'
    return (f'IF({x}="","",IF(ISNUMBER({x}),IF({x}>19000000,DATE(INT({x}/10000),MOD(INT({x}/100),100),MOD({x},100)),'
            f'IF({x}<36526,"",INT({x}))),'
            f'IFERROR(IF(ISERROR(FIND("-",{dp})),DATE(LEFT({dp},4),MID({dp},5,2),RIGHT({dp},2)),'
            f'DATE(LEFT({dp},4),MID({dp},6,FIND("-",{dp},6)-6),MID({dp},FIND("-",{dp},6)+1,2))),"")))')


def num(x):
    """粘贴来的金额可能是文本、带 ¥ 或千分位逗号"""
    return f'IF({x}="",0,IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"¥",""),"￥",""),",",""),0))'
