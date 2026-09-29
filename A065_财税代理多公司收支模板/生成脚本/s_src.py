# -*- coding: utf-8 -*-
"""【流水1～流水8】每个导入账户一张：网银 / 微信 / 支付宝导出的文件整份原样粘进来（连标题、表头一起也行），
   按表头文字自动认出「日期、收入、支出、余额、对方户名、摘要……」各在第几列，认出来的每一笔自动进【资金台帐】。
   【手工记账】现金、个人代付等没有导出文件的账户，一笔一行手工记。
   认不出来的往来单位 / 收支项目，就在这一笔右边的 V、W 列手工选（跟着这一行走，不会错位）。"""
from openpyxl.formatting.rule import FormulaRule
from common import *

# 表头关键字（按顺序找，先找到的算；* 是通配符）
PAT = {
    '日期': ['交易日期', '记账日期', '入账日期', '日期', '*交易日期*', '*记账日期*', '*日期*'],
    '时间': ['交易时间', '时间', '*交易时间*', '*时间*'],
    '收入': ['收入金额', '收入', '贷方发生额', '贷方金额', '贷方', '存入', '*收入金额*', '*贷方发生额*', '*贷方金额*',
           '收入(*', '收入（*'],
    '支出': ['支出金额', '支出', '借方发生额', '借方金额', '借方', '取出', '*支出金额*', '*借方发生额*', '*借方金额*',
           '支出(*', '支出（*'],
    '单列金额': ['金额', '金额(元)', '金额（元）', '交易金额', '发生额', '*交易金额*', '*金额*', '*发生额*'],
    '收支标志': ['收/支', '收支', '收支类型', '借贷标志', '借/贷', '借贷', '借贷方向', '*收/支*', '*借贷*', '*收支*'],
    '余额': ['余额', '账户余额', '当前余额', '*余额*'],
    '对方账号': ['对方账号', '对方账户', '对方卡号', '对方帐号', '*对方*账号*', '*对方*帐号*', '*对方*卡号*'],
    '对方户名': ['对方户名', '对方名称', '对方单位', '对方账户名', '对方账户名称', '对方单位名称', '交易对方', '收/付款方',
             '*对方户名*', '*交易对方*', '*对方*单位*', '*对方*名称*'],
    '开户行': ['对方开户行', '对方行名', '对方开户机构', '对方开户银行', '对方银行', '*对方*开户*', '*对方*行名*'],
    '摘要': ['摘要', '交易摘要', '交易类型', '交易分类', '*摘要*'],
    '用途/备注': ['用途', '附言', '交易附言', '商品', '商品说明', '备注', '交易备注', '*用途*', '*附言*', '*商品*'],
    '支付方式': ['支付方式', '收/付款方式', '*付款方式*', '*支付方式*'],
    '状态': ['当前状态', '交易状态', '*状态*'],
    '单号': ['交易单号', '交易订单号', '交易流水号', '流水号', '凭证号', '*单号*', '*流水号*'],
}
# 手工记账的列（固定）
MAN_COLS = ['账户', '日期', '收入金额', '支出金额', '对方户名', '摘要', '往来单位', '收支项目', '备注', '对方账号', '银行余额']
MAN_MAP = {'账户列': 1, '日期': 2, '收入': 3, '支出': 4, '对方户名': 5, '摘要': 6, '往来单位(手工)': 7,
           '收支项目(手工)': 8, '备注(手工)': 9, '对方账号': 10, '余额': 11}
C_SRC = 'FF2F75B5'


def _find(hrow, pats):
    """在表头那一行里按顺序找，返回列号；都没有返回 0"""
    f = '0'
    for p in reversed(pats):
        f = f'IFERROR(MATCH("{p}",{hrow},0),{f})'
    return f


def mapping_letters():
    return {n: CL(i + 1) for i, n in enumerate(FIELDS)}      # 第 4～7 行：字段在 A～S


def build_src(wb, ctx, j):
    """j：1～N_IMP 导入账户"""
    name = SRC_SHEETS[j - 1]
    ws = wb.create_sheet(name)
    acc_row = AC_R0 + j - 1
    acc = f'{SH_BASE}!$I${acc_row}'
    widths(ws, {**{c: 13 for c in S_RAW}, S_STAT: 30, S_MP: 18, S_MN: 14, S_MI: 12})
    ws.column_dimensions['A'].width = 18
    ws.merge_cells(f'A1:{S_MN}1')
    put(ws, 'A1', f'="流 水 {j} · "&IF({acc}="","（基础资料 ② 第 {j} 个账户还没填）",{acc}&"（"&{SH_BASE}!$J${acc_row}&"）")',
        F_TITLE, fill(C_SRC), align=AC, border=False)
    ws.row_dimensions[1].height = 33
    ws.merge_cells(f'A2:{S_MN}2')
    put(ws, 'A2', f'💡 把这个账户导出的流水文件（Excel）打开，全选复制，粘到 A{S_R0}（第一次连表头一起粘，会自动认出哪列是日期、收入、支出、余额、对方户名、摘要）；'
                  f'以后每次导出的新流水粘到下面第一个空行，表头、标题、合计行粘进来也没关系，会自动跳过。'
                  f'认出来的每一笔自动进【资金台帐】，U 列显示结果；✗ 的就在同一行 V 列选往来单位、W 列选收支项目（跟着这一笔走）。'
                  f'第 5 行是自动认的列号，认错了在第 6 行填正确的列号（A 列＝1，B 列＝2……）。'
                  f'微信、支付宝里用银行卡付的钱（银行流水里已经有）、交易关闭的、零钱通来回倒的会自动跳过，免得记两次。',
        F_TIP, FILL_TIP, align=ALW, border=False)
    ws.row_dimensions[2].height = 64
    # 第 3 行：状态条
    put(ws, 'A3', '账户', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B3', f'=IF({acc}="","（没填）",{acc})', F_SEL, FILL_SEL, align=AC)
    put(ws, 'C3', '所属公司', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'D3', f'={SH_BASE}!$J${acc_row}&""', F_SEL, fill('FFF2F2F2'), align=AC)
    put(ws, 'E3', '已认出', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'F3', f'=${S_CUM}${S_R1}&" 笔"', F_KPI_V, fill('FFF2F2F2'), align=AC)
    put(ws, 'G3', '格式', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells(f'H3:{S_MN}3')
    L = mapping_letters()
    e = lambda n: f'${L[n]}${S_EFF}'
    show = lambda n: f'IF({e(n)}>0,"{n}="&CHAR(64+{e(n)})&" ","")'
    put(ws, 'H3', f'=IF(${S_HROW}${S_EFF}=0,IF(COUNTA($A${S_R0}:$T${S_R0 + S_HSCAN - 1})=0,'
                  f'"还没有流水：第一次导出后整份粘到 A{S_R0}（连表头），会自动认格式",'
                  f'"✗ 前 {S_HSCAN} 行里没找到表头（要有「日期/时间」和「金额/收入/支出」那一行）：在 U6 填表头在第几行，或者第 6 行直接填列号"),'
                  f'"表头在第 "&${S_HROW}${S_EFF}&" 行 → "&{show("日期")}&{show("收入")}&{show("支出")}&{show("单列金额")}&{show("收支标志")}'
                  f'&{show("余额")}&{show("对方户名")}&{show("摘要")}&IF(AND({e("日期")}>0,OR(AND({e("收入")}>0,{e("支出")}>0),{e("单列金额")}>0)),"✓","✗ 缺日期或金额列"))',
        F_AUTOB, fill('FFF2F2F2'), align=AL)
    for c in range(CI('I'), CI(S_MN) + 1):
        ws.cell(3, c).border = BD
    ws.row_dimensions[3].height = 24
    # 第 4～7 行：字段列号
    for n, col in L.items():
        put(ws, f'{col}{S_LBL}', n, F_HDR, fill('FF7F7F7F'), align=ACW)
    put(ws, f'{S_HROW}{S_LBL}', '表头在第几行', F_HDR, fill('FF7F7F7F'), align=ACW)
    labels = {S_AUTO: '自动认的', S_MAN: '手工改（空＝用自动）', S_EFF: '实际用'}
    for r, t in labels.items():
        put(ws, f'V{r}', t, F_NOTE, align=AL, border=False)
    ws.row_dimensions[S_LBL].height = 30
    # 表头行：前 S_HSCAN 行里第一行同时有「日期/时间」和「金额/收入/支出…」的
    hscan = f'${S_HF}${S_R0}:${S_HF}${S_R0 + S_HSCAN - 1}'
    put(ws, f'{S_HROW}{S_AUTO}', f'=IFERROR(MATCH(1,{hscan},0)+{S_R0 - 1},0)', F_AUTO, FILL_AUTO, align=AC)
    put(ws, f'{S_HROW}{S_MAN}', None, F_IN, FILL_IN, align=AC)
    put(ws, f'{S_HROW}{S_EFF}', f'=IF(N({S_HROW}{S_MAN})>0,{S_HROW}{S_MAN},{S_HROW}{S_AUTO})', F_AUTOB, FILL_AUTO, align=AC)
    hrow = f'INDEX($A$1:$T${S_R1},${S_HROW}${S_EFF},0)'
    for n, col in L.items():
        k = FI[n]
        if k <= N_DET:
            if n == '日期':
                f = f'IF(${S_HROW}${S_EFF}=0,0,IF({_find(hrow, PAT["日期"])}>0,{_find(hrow, PAT["日期"])},{_find(hrow, PAT["时间"])}))'
            elif n == '时间':
                t = _find(hrow, PAT['时间'])
                f = f'IF(${S_HROW}${S_EFF}=0,0,IF(AND({_find(hrow, PAT["日期"])}>0,{t}>0),IF({t}<>{_find(hrow, PAT["日期"])},{t},0),0))'
            elif n == '单列金额':
                f = (f'IF(OR(${S_HROW}${S_EFF}=0,AND(${L["收入"]}${S_AUTO}>0,${L["支出"]}${S_AUTO}>0)),0,'
                     f'{_find(hrow, PAT[n])})')
            elif n == '收支标志':
                f = f'IF(OR(${S_HROW}${S_EFF}=0,${L["单列金额"]}${S_AUTO}=0),0,{_find(hrow, PAT[n])})'
            else:
                f = f'IF(${S_HROW}${S_EFF}=0,0,{_find(hrow, PAT[n])})'
            put(ws, f'{col}{S_AUTO}', '=' + f, F_AUTO, FILL_AUTO, align=AC)
            put(ws, f'{col}{S_MAN}', None, F_IN, FILL_IN, align=AC)
            put(ws, f'{col}{S_EFF}', f'=IF(N({col}{S_MAN})>0,{col}{S_MAN},{col}{S_AUTO})', F_AUTOB, FILL_AUTO, align=AC)
        else:
            v = {'往来单位(手工)': CI(S_MP), '收支项目(手工)': CI(S_MI), '备注(手工)': CI(S_MN), '账户列': 0}[n]
            put(ws, f'{col}{S_AUTO}', v, F_AUTO, FILL_AUTO, align=AC)
            put(ws, f'{col}{S_MAN}', None, fill_=fill('FFF2F2F2'))
            put(ws, f'{col}{S_EFF}', f'={col}{S_AUTO}', F_AUTOB, FILL_AUTO, align=AC)
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='20', allow_blank=True,
                        showErrorMessage=True, errorTitle='列号', error='填 1～20 的数字（A 列＝1）')
    ws.add_data_validation(dv)
    dv.add(f'A{S_MAN}:O{S_MAN}')
    dv.add(f'{S_HROW}{S_MAN}')
    _data_area(ws, ctx, j, name, e)
    return ws


def _data_area(ws, ctx, j, name, e, manual=False):
    # 第 8 行：列标题
    for i, c in enumerate(S_RAW):
        put(ws, f'{c}{S_HDR}', MAN_COLS[i] if manual and i < len(MAN_COLS) else f'第{i + 1}列',
            F_HDR if manual and i < len(MAN_COLS) else F_NOTE, fill('FF5B9BD5') if manual and i < len(MAN_COLS) else fill('FFDDEBF7'),
            align=ACW)
    for c, t, colr in ((S_STAT, '自动认的结果\n（进资金台帐没有）', 'FF7F7F7F'), (S_MP, '往来单位\n（认不出/认错了才选）', 'FFBF8F00'),
                       (S_MI, '收支项目\n（手工改）', 'FFBF8F00'), (S_MN, '备注', 'FFBF8F00')):
        if manual and c in (S_MP, S_MI, S_MN):
            continue
        put(ws, f'{c}{S_HDR}', t, F_HDR, fill(colr), align=ACW)
    ws.row_dimensions[S_HDR].height = 32
    for c, t in ((S_HF, '表头?'), (S_OK, '流水行?'), (S_CUM, '累计')):
        ws[f'{c}{S_HDR}'] = t
        ws[f'{c}{S_HDR}'].font = F_HELP
    ws[f'{S_CUM}{S_HDR}'] = 0                                # 累计从 0 起（资金台帐二分查找用）
    off = f'{SH_AUX}!${AUX_OFF}${AUX_R0 + j - 1}'
    dcol = e('日期')
    for r in range(S_R0, S_R1 + 1):
        row = f'$A{r}:$T{r}'
        x = f'INDEX({row},1,{dcol})'
        if r < S_R0 + S_HSCAN and not manual:
            ws[f'{S_HF}{r}'] = (f'=IF(COUNTA({row})<2,0,IF(AND(COUNTIF({row},"*日期*")+COUNTIF({row},"*时间*")>0,'
                                f'COUNTIF({row},"*金额*")+COUNTIF({row},"*收入*")+COUNTIF({row},"*支出*")+COUNTIF({row},"*发生额*")'
                                f'+COUNTIF({row},"*收/支*")+COUNTIF({row},"*借贷*")>0),1,0))')
        # 是不是一笔流水：日期那一格像日期（真日期，或 20 开头的文字）；手工记账还要选了账户
        ok = (f'IF({dcol}=0,0,IF(OR(AND(ISNUMBER({x}),{x}>36526),AND(ISTEXT({x}),LEFT(TRIM({x}),2)="20",LEN(TRIM({x}))>=8)),1,0))')
        if manual:
            ok = f'IF($A{r}="",0,{ok})'
        ws[f'{S_OK}{r}'] = '=' + ok
        ws[f'{S_CUM}{r}'] = f'={S_CUM}{r - 1}+{S_OK}{r}'
        ws[f'{S_STAT}{r}'] = (f'=IF({S_OK}{r}=1,IFERROR(INDEX({SH_CASH}!${J_CHK}${J_R0}:${J_CHK}${J_R1},{off}+{S_CUM}{r}),""),'
                              + (f'IF(N({S_HF}{r})=1,"（表头）",""))' if (r < S_R0 + S_HSCAN and not manual) else '"")'))
        for c in (S_HF, S_OK, S_CUM):
            ws[f'{c}{r}'].font = F_HELP
    style_rows(ws, S_R0, S_R1, S_RAW, fills={c: FILL_PASTE for c in S_RAW},
               fmts={c: '@' for c in S_RAW} if False else None)
    style_rows(ws, S_R0, S_R1, [S_STAT], auto=[S_STAT], aligns={S_STAT: AL})
    if not manual:
        style_rows(ws, S_R0, S_R1, [S_MP, S_MI, S_MN], fills={S_MP: FILL_IN, S_MI: FILL_IN, S_MN: FILL_IN},
                   aligns={S_MP: AL, S_MN: AL})
        dv_list(ws, f'{S_MP}{S_R0}:{S_MP}{S_R1}', f'={SH_AUX}!$C$1:$C$620', '认错了或认不出才选；也可以直接打名字', stop=False)
        dv_list(ws, f'{S_MI}{S_R0}:{S_MI}{S_R1}', f'={IT_NAMES}', '自动认的不对才选：比如社保挂靠、续费改代办')
    rng = f'{S_STAT}{S_R0}:{S_STAT}{S_R1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${S_STAT}{S_R0},1)="✗"'], fill=FILL_WARN,
                                                   font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${S_STAT}{S_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(${S_STAT}{S_R0},1)="√"'],
                                                   font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    hide(ws, S_HF, 'Z', S_OK, S_CUM)
    ws.freeze_panes = f'A{S_R0}'


def build_man(wb, ctx):
    ws = wb.create_sheet(SH_MAN)
    j = N_SRC
    widths(ws, {'A': 13, 'B': 12, 'C': 12, 'D': 12, 'E': 26, 'F': 22, 'G': 22, 'H': 12, 'I': 16, 'J': 20, 'K': 13,
                **{c: 4 for c in S_RAW[11:]}, S_STAT: 30})
    ws.merge_cells(f'A1:{S_STAT}1')
    put(ws, 'A1', '手 工 记 账（现金、个人代付等没有导出文件的账户）', F_TITLE, fill(C_SRC), align=AC, border=False)
    ws.row_dimensions[1].height = 33
    ws.merge_cells(f'A2:{S_STAT}2')
    put(ws, 'A2', '💡 一笔一行：A 选账户（在【基础资料】② 登记的，一般是现金、个人卡），B 日期，C/D 收入或支出，E 对方户名，F 摘要；'
                  'G 往来单位、H 收支项目认得出就不用填，认不出、要改（社保挂靠、费用……）才填。记好的自动进【资金台帐】，U 列显示结果。'
                  '银行、微信、支付宝有导出文件的账户不要在这里记，粘到它自己的【流水】表。',
        F_TIP, FILL_TIP, align=ALW, border=False)
    ws.row_dimensions[2].height = 40
    # 第 4～7 行：固定列号（隐藏）
    L = mapping_letters()
    for n, col in L.items():
        ws[f'{col}{S_LBL}'] = n
        ws[f'{col}{S_EFF}'] = MAN_MAP.get(n, 0)
        ws[f'{col}{S_LBL}'].font = ws[f'{col}{S_EFF}'].font = F_HELP
    ws[f'{S_HROW}{S_EFF}'] = S_HDR
    for r in range(3, S_EFF + 1):
        ws.row_dimensions[r].hidden = True
    e = lambda n: f'${L[n]}${S_EFF}'
    _data_area(ws, ctx, j, SH_MAN, e, manual=True)
    fm = {'B': DATE, 'C': MONEY, 'D': MONEY, 'K': MONEY, 'J': '@'}
    for r in range(S_R0, S_R1 + 1):
        for c, f in fm.items():
            ws[f'{c}{r}'].number_format = f
        for c in 'GHI':
            ws[f'{c}{r}'].fill = FILL_IN
    dv_list(ws, f'A{S_R0}:A{S_R1}', f'={AC_NAMES}', '哪个账户（在【基础资料】② 登记）')
    dv_date(ws, f'B{S_R0}:B{S_R1}')
    dv_list(ws, f'G{S_R0}:G{S_R1}', f'={SH_AUX}!$C$1:$C$620', '认不出才选；也可以直接打名字', stop=False)
    dv_list(ws, f'H{S_R0}:H{S_R1}', f'={IT_NAMES}', '认不出、要改才选')
    for c in S_RAW[11:]:
        ws.column_dimensions[c].hidden = True
    return ws


def define_names(wb):
    """资金台帐按「第几张流水表」取数：流水区1…、累计区1… 是各张表的整片区域和累计列
       （名字不能起成 SRC1 这种样子——Excel 会当成单元格地址）"""
    from openpyxl.workbook.defined_name import DefinedName
    for j, nm in enumerate(SRC_SHEETS, 1):
        for key, ref in ((f'流水区{j}', f"'{nm}'!$A$1:${S_MN}${S_R1}"), (f'累计区{j}', f"'{nm}'!${S_CUM}${S_HDR}:${S_CUM}${S_R1}")):
            wb.defined_names[key] = DefinedName(key, attr_text=ref)


def build_all(wb, ctx):
    for j in range(1, N_IMP + 1):
        build_src(wb, ctx, j)
    build_man(wb, ctx)
    define_names(wb)
    # 演示：流水1 放你发来的农行导出文件原样
    ws = wb[SRC_SHEETS[0]]
    for i, row in enumerate(ctx['bank_raw']):
        for k, v in enumerate(row):
            c = ws.cell(row=S_R0 + i, column=k + 1, value=v)
            if isinstance(v, str) and v.isdigit():
                c.number_format = '@'
