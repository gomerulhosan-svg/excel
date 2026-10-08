# -*- coding: utf-8 -*-
"""10/08 这一轮：以你 10 月 5 日的版本（../参考/1008上传/）为底，改四件事。

  ① 67 号矿井溢价款（4 笔共 450 万人民币）改记 其他应收款-保证金（新科目 122106）
  ④ 退回来的溢价款（收 4 笔 / 收 3 笔人民币）也改记 其他应收款-保证金
     —— 两条都是在【现金流水总表】「指定对方科目」填 122106；这两类的现金流量归「其他与经营活动有关的现金」。
  ③ 凭证按天合并：同一天、同一个资金账户、同一方向（收 / 付）的流水并成一张凭证——
     收款＝一借多贷（借 本账户 合计一行，贷 各对方科目），付款＝多借一贷（借 各对方科目，贷 本账户 合计一行）；
     苏姆现金、人民币现金、美元现金、各银行账户各自分开。填了「凭证合并号」的照旧按合并号并。
  ② 【记账凭证】打印：A4 纵向，一张纸上下两页凭证，中间留宽便于裁开分别装订；
     月末结转凭证（折旧、调汇、成本、销售成本、兑换差额，加上新做的「结转损益」）和
     12 月 31 日的年末结转（本年利润 → 利润分配-未分配利润）都进凭证汇总、按号打印。
     结转损益以后损益科目月末余额为 0，利润表改取「不含结转损益」的发生额，数不变。

跑法：见 build_1008.py（这里只提供 apply(wb)）。
"""
import re
from copy import copy

from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break, RowBreak

# ---------------- 尺寸（跟你那本一致） ----------------
C0, C1 = 6, 3965            # 现金流水总表 数据行
JE0, JE1 = 4, 16063         # 记账分录 原有槽位
FLOW_K = 4000               # 记账分录 流水块：每笔流水 2 行，共 4000 笔
IX1 = 6863                  # _凭证索引 原有最后一行
NM = 32                     # 月份个数（起始月份起 32 个月）
NPL = 200                   # 损益科目槽位（科目表里 5 开头的科目，现有约 120 个）
NJ = 60                     # 每张结转损益凭证最多几行损益科目
NY = 3                      # 年末结转：3 个年度
MBLK = NJ + 1               # 每月一块：60 行损益 + 1 行本年利润
CLOSE0 = JE1 + 1                        # 16064 结转损益块开头
YE0 = CLOSE0 + NM * MBLK                # 18016 年末结转块开头
JE_END = YE0 + NY * 2 - 1               # 18021 记账分录新的最后一行
IXM0 = IX1 + 1                          # 6864 _凭证索引：结转损益 32 行
IXY0 = IXM0 + NM                        # 6896 _凭证索引：年末结转 3 行
IX_END = IXY0 + NY - 1                  # 6898
ID_M, ID_Y = 1200000, 1300000
SAVE_CODE = '122106'
SAVE_ROWS = {6: '170101', 7: '170101', 8: '170101', 61: '170101', 1631: '170101', 2239: '170101'}

# _结转损益 辅助表布局
H_M0 = 7                                 # G 列＝第 1 个月
H_CUM0 = 40                              # AN 列＝第 1 个月的累计条数
R_TOT, R_CNT, R_YR = 206, 207, 208
R_K0, R_A0, R_C0 = 210, 272, 334         # 压缩后：第 j 行用哪个科目槽位 / 金额 / 科目编码
R_PL = 395                               # 本年利润（借-贷）
R_Y0 = 398                               # 年末结转表 398..400

CASH = '现金流水总表'
JE = '记账分录'
IX = '_凭证索引'
HP = '_结转损益'


def mcol(m):            # 第 m 个月在辅助表的列字母
    return CL(H_M0 + m - 1)


def cumcol(m):
    return CL(H_CUM0 + m - 1)


def cp(src, dst):
    dst.font = copy(src.font)
    dst.fill = copy(src.fill)
    dst.border = copy(src.border)
    dst.alignment = copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy(src.protection)


def expect(ws, ref, want, bad):
    got = ws[ref].value
    if got != want:
        bad.append(f'{ws.title}!{ref}\n   有：{got}\n   该：{want}')


def cx(col, k):          # 记账分录里取现金流水某列第 K 笔
    return f'INDEX({CASH}!${col}${C0}:${col}${C1},$K{k})'


# =====================================================================
# ① ④ 保证金科目、6 笔流水改科目
# =====================================================================
def step_deposit(wb, log):
    ws = wb['科目表']
    r = 344
    assert ws[f'A{r}'].value is None and ws['A343'].value == '6000', '科目表第 344 行不是空行'
    assert all(str(ws.cell(i, 1).value) != SAVE_CODE for i in range(5, 455))
    for col, v in zip('ABCDEFGHI', [SAVE_CODE, '保证金', '流动资产', '借', None, 'SOM/RMB/USD', '否', '启用', '否']):
        c = ws[f'{col}{r}']
        cp(ws[f'{col}40'], c)
        c.value = v
    ws[f'A{r}'].number_format = '@'
    log.append('科目表第 344 行加 122106 其他应收款-保证金')

    cs = wb[CASH]
    for row, old in SAVE_ROWS.items():
        assert str(cs[f'V{row}'].value) == old, (row, cs[f'V{row}'].value)
        assert '67号矿井' in str(cs[f'G{row}'].value) or str(cs[f'G{row}'].value).startswith('收'), cs[f'G{row}'].value
        cs[f'V{row}'].value = SAVE_CODE
    log.append(f'现金流水 {sorted(SAVE_ROWS)} 行「指定对方科目」170101 → 122106')

    # 现金流量：122106 归「收到/支付其他与经营活动有关的现金」（2 / 6）
    rule = ('IFERROR(IF($O{r}>0,N(INDEX(记账规则!$G$5:$G$64,MATCH($E{r},类别表,0))),'
            'N(INDEX(记账规则!$H$5:$H$64,MATCH($E{r},类别表,0)))),0)')
    head = ('=IF($AE{r}=1,IF($V{r}="",{R},IF(OR(LEFT($AB{r},4)="1001",LEFT($AB{r},4)="1002",LEFT($AB{r},4)="1012",'
            '$AB{r}="122104",$AB{r}="122105"),0,')
    tail = ('IF(OR(LEFT($AB{r},2)="16",LEFT($AB{r},2)="17"),IF($O{r}>0,10,12),IF(AND(LEFT($AB{r},1)="5",{R}=0),'
            'IF($O{r}>0,2,6),{R}))')
    bad = []
    for r in range(C0, C1 + 1):
        R = rule.format(r=r)
        old = (head + tail + ')),"")').format(r=r, R=R)
        new = (head + 'IF($AB{r}="' + SAVE_CODE + '",IF($O{r}>0,2,6),' + tail + '))),"")').format(r=r, R=R)
        expect(cs, f'AF{r}', old, bad)
        cs[f'AF{r}'].value = new
    assert not bad, '\n'.join(bad[:5])
    log.append('现金流水 AF（现金流量项目）：对方科目 122106 归经营活动「其他」')


# =====================================================================
# ③ 凭证按天、按账户、按收付合并
# =====================================================================
def step_merge(wb, log):
    cs = wb[CASH]
    bad = []
    for r in range(C0, C1 + 1):
        expect(cs, f'AG{r}', f'=IF(OR($AE{r}=0,$X{r}=""),"",COUNTIFS($X$6:$X{r},$X{r},$AE$6:$AE{r},1))', bad)
        expect(cs, f'AH{r}', f'=IF($AG{r}=1,$X{r}&"","")', bad)
        expect(cs, f'AI{r}', f'=IF($AE{r}=0,"",IF($X{r}="",100000+ROW()-5,100000+IFERROR(MATCH($X{r}&"",$AH$6:$AH$3965,0),ROW()-5)))', bad)
        for col in ('AK', 'AL', 'AM', 'AN'):
            if cs[f'{col}{r}'].value is not None:
                bad.append(f'{col}{r} 不是空的')
    assert not bad, '\n'.join(bad[:5])

    # 新加 4 个辅助列（隐藏）
    # AH 原来是「组首」（只给旧的 AI 用），改成「查找用的合并键」：把 ~ * ? 转义，账户名、合并号里有这些字符也不会被当通配符
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('=') and (f'{CASH}!$AH' in v or (ws.title == CASH and '$AH$' in v and c.column_letter != 'AI')):
                    raise AssertionError(f'{ws.title}!{c.coordinate} 还在用 现金流水总表 AH：{v[:120]}')
    cs['AH5'].value = '合并键\n(查找用)'
    hdr = {'AK': '合并键', 'AL': '组笔数', 'AM': '组合计\n(苏姆)', 'AN': '合并行摘要'}
    for col, t in hdr.items():
        cp(cs['AJ5'], cs[f'{col}5'])
        cs[f'{col}5'].value = t
        cs.column_dimensions[col].width = 12
        cs.column_dimensions[col].hidden = True
    rng = lambda c: f'${c}${C0}:${c}${C1}'
    for r in range(C0, C1 + 1):
        for col in ('AK', 'AL', 'AM', 'AN'):
            cp(cs[f'AJ{r}'], cs[f'{col}{r}'])
        cs[f'AM{r}'].number_format = '#,##0.00'
        cs[f'AK{r}'].value = (f'=IF($AE{r}=0,"",IF($X{r}<>"","X|"&$X{r},INT($A{r})&"|"&$B{r}&"|"&IF($O{r}>0,"收","支")))')
        cs[f'AH{r}'].value = f'=IF($AK{r}="","",SUBSTITUTE(SUBSTITUTE(SUBSTITUTE($AK{r},"~","~~"),"*","~*"),"?","~?"))'
        cs[f'AG{r}'].value = f'=IF($AE{r}=0,"",COUNTIF($AK${C0}:$AK{r},$AH{r}))'
        cs[f'AI{r}'].value = f'=IF($AE{r}=0,"",100000+MATCH($AH{r},{rng("AK")},0))'
        cs[f'AL{r}'].value = f'=IF($AG{r}=1,COUNTIF({rng("AK")},$AH{r}),"")'
        cs[f'AM{r}'].value = f'=IF(AND($AG{r}=1,$X{r}=""),ABS(SUMIF({rng("AK")},$AH{r},{rng("O")})),"")'
        cs[f'AN{r}'].value = (f'=IF(OR($AG{r}<>1,$X{r}<>""),"",IF($AL{r}=1,$AJ{r},$B{r}&" "&MONTH($A{r})&"月"&DAY($A{r})&"日"'
                              f'&IF($O{r}>0,"收款","付款")&$AL{r}&"笔合计"&IF($D{r}=本位币,"","（"&$D{r}&" "'
                              f'&FIXED(ABS(SUMIF({rng("AK")},$AH{r},{rng("N")})),2)&"）")))')
    a2 = cs['A2'].value
    s1, s2 = '每一笔自动生成一张记账凭证。', '几笔想并成一张凭证：「凭证合并号」写同一个号（要在同一个月）。'
    assert s1 in a2 and s2 in a2
    cs['A2'].value = a2.replace(s1, '凭证自动按天并：同一天、同一个账户的收款并成一张（一借多贷），付款并成一张（多借一贷），各账户分开。').replace(
        s2, '想把别的几笔并成一张：「凭证合并号」写同一个号（要在同一个月；填了合并号的就按合并号并）。')
    log.append('现金流水 AG/AH/AI 改按「日期＋账户＋收付」分组，新加 AK~AN 辅助列（隐藏）')

    je = wb[JE]
    bad = []
    for k in range(1, FLOW_K + 1):
        a, b = JE0 + 2 * (k - 1), JE0 + 2 * (k - 1) + 1
        assert je[f'K{a}'].value == k and je[f'K{b}'].value == k and je[f'J{a}'].value == '流水'
        expect(je, f'C{a}', f'=IF($L{a}="","",{cx("AJ", a)})', bad)
        expect(je, f'F{a}', f'=IF($L{a}="",0,{cx("AC", a)}+{cx("AD", a)})', bad)
        expect(je, f'G{a}', 0, bad)
        expect(je, f'H{a}', f'=IF($L{a}="","",{cx("S", a)}&"")', bad)
        expect(je, f'I{a}', f'=IF($L{a}="","",{cx("R", a)}&"")', bad)
        expect(je, f'M{a}', f'=IF($P{a}=0,"",$Q${a}+SUM($P${a}:$P{a}))', bad)
        expect(je, f'Q{a}', f'=IF($L{a}="",0,(IF({cx("AG", a)}="",1,{cx("AG", a)})-1)*2)', bad)
        expect(je, f'C{b}', f'=$C{a}', bad)
        expect(je, f'F{b}', 0, bad)
        expect(je, f'G{b}', f'=$F{a}', bad)
        expect(je, f'H{b}', f'=$H{a}', bad)
        expect(je, f'I{b}', f'=$I{a}', bad)
        expect(je, f'M{b}', f'=IF($P{b}=0,"",$Q${a}+SUM($P${a}:$P{b}))', bad)
        if bad:
            break
    assert not bad, '\n'.join(bad[:5])

    for k in range(1, FLOW_K + 1):
        a, b = JE0 + 2 * (k - 1), JE0 + 2 * (k - 1) + 1
        X, AC, AD, AG, AL = (lambda r: (cx('X', r), cx('AC', r), cx('AD', r), cx('AG', r), cx('AL', r)))(a)
        # 第 1 行＝借方那一行：收款时是「本账户（并成合计一行，只在组里第一笔上）」，付款时是「对方科目（每笔一行）」
        je[f'C{a}'].value = f'=IF($L{a}="","",IF(AND({X}="",{AC}>0),{cx("AN", a)},{cx("AJ", a)}))'
        je[f'F{a}'].value = f'=IF($L{a}="",0,IF(OR({X}<>"",{AC}=0),{AC}+{AD},IF({AG}=1,{cx("AM", a)},0)))'
        je[f'H{a}'].value = f'=IF($L{a}="","",IF(AND({X}="",{AC}>0,N({AL})<>1),"",{cx("S", a)}&""))'
        je[f'I{a}'].value = f'=IF($L{a}="","",IF(AND({X}="",{AC}>0,N({AL})<>1),"",{cx("R", a)}&""))'
        je[f'Q{a}'].value = f'=IF($L{a}="",0,IF({X}<>"",({AG}-1)*2+1,IF({AC}>0,1,{AG})))'
        je[f'M{a}'].value = f'=IF($P{a}=0,"",$Q{a})'
        # 第 2 行＝贷方那一行：收款时是「对方科目（每笔一行）」，付款时是「本账户（合计一行）」
        X, AC, AD, AG, AL = (lambda r: (cx('X', r), cx('AC', r), cx('AD', r), cx('AG', r), cx('AL', r)))(b)
        je[f'C{b}'].value = f'=IF($L{b}="","",IF(AND({X}="",{AD}>0),{cx("AN", b)},{cx("AJ", b)}))'
        je[f'G{b}'].value = f'=IF($L{b}="",0,IF(OR({X}<>"",{AC}>0),{AC}+{AD},IF({AG}=1,{cx("AM", b)},0)))'
        je[f'H{b}'].value = f'=IF($L{b}="","",IF(AND({X}="",{AD}>0,N({AL})<>1),"",{cx("S", b)}&""))'
        je[f'I{b}'].value = f'=IF($L{b}="","",IF(AND({X}="",{AD}>0,N({AL})<>1),"",{cx("R", b)}&""))'
        je[f'Q{b}'].value = f'=IF($L{b}="",0,IF({X}<>"",({AG}-1)*2+2,IF({AC}>0,{AG}+1,N({AL})+1)))'
        cp(je[f'Q{a}'], je[f'Q{b}'])
        je[f'M{b}'].value = f'=IF($P{b}=0,"",$Q{b})'
    je['Q3'].value = '条号'
    log.append('记账分录 流水块：收款借方 / 付款贷方 的本账户并成合计一行，条号按组重排')

    ix = wb[IX]
    ci = lambda col, k: f'INDEX({CASH}!${col}${C0}:${col}${C1},{k})'
    bad = []
    for r in range(4, 4 + FLOW_K):
        k = r - 3
        assert ix[f'A{r}'].value == 100000 + k and ix[f'B{r}'].value == '流水'
        expect(ix, f'F{r}', f'=IF($C{r}=0,0,IF({ci("X", k)}="",2,2*COUNTIF({CASH}!$AI${C0}:$AI${C1},$A{r})))', bad)
        expect(ix, f'J{r}', f'=IF($C{r}=0,"",{ci("AJ", k)})', bad)
        expect(ix, f'K{r}', f'=IF($C{r}=0,"",IF({ci("X", k)}="",N({ci("Y", k)}),SUMIF({CASH}!$AI${C0}:$AI${C1},$A{r},{CASH}!$Y${C0}:$Y${C1})))', bad)
        if bad:
            break
    assert not bad, '\n'.join(bad[:5])
    for r in range(4, 4 + FLOW_K):
        k = r - 3
        ix[f'F{r}'].value = f'=IF($C{r}=0,0,IF({ci("X", k)}="",N({ci("AL", k)})+1,2*N({ci("AL", k)})))'
        ix[f'J{r}'].value = f'=IF($C{r}=0,"",IF(AND({ci("X", k)}="",{ci("AC", k)}>0),{ci("AN", k)},{ci("AJ", k)}))'
        ix[f'K{r}'].value = f'=IF($C{r}=0,"",SUMIF({CASH}!$AI${C0}:$AI${C1},$A{r},{CASH}!$Y${C0}:$Y${C1}))'
    log.append('_凭证索引 流水凭证：条数＝笔数＋1、摘要取合并行、附单据按组合计')


# =====================================================================
# ② 结转损益（每月）+ 年末结转本年利润
# =====================================================================
def step_close(wb, log):
    je = wb[JE]
    # 1) 记账分录 R、S 两个辅助列：损益科目的「槽位×100＋月序」和「借-贷」
    for col, t in (('R', '损益键'), ('S', '损益\n借-贷')):
        cp(je['Q3'], je[f'{col}3'])
        je[f'{col}3'].value = t
        je.column_dimensions[col].width = 8
        je.column_dimensions[col].hidden = True
    for r in range(JE0, JE1 + 1):
        assert je[f'R{r}'].value is None and je[f'S{r}'].value is None
        je[f'R{r}'].value = (f'=IF($L{r}="","",IF(LEFT($D{r},1)<>"5","",IFERROR(MATCH($D{r},{HP}!$C$5:$C${4 + NPL},0)*100'
                             f'+(YEAR($A{r})-YEAR(起始月份))*12+MONTH($A{r})-MONTH(起始月份)+1,"")))')
        je[f'S{r}'].value = f'=IF($R{r}="",0,$F{r}-$G{r})'
        cp(je[f'P{r}'], je[f'R{r}'])
        cp(je[f'P{r}'], je[f'S{r}'])

    # 2) 辅助表 _结转损益
    assert HP not in wb.sheetnames
    hp = wb.create_sheet(HP)
    hp.sheet_state = 'hidden'
    f9 = Font(name='微软雅黑', sz=9)
    fb = Font(name='微软雅黑', sz=9, bold=True)
    hp['A1'] = '结转损益辅助（自动，不用动）：每月把损益科目（科目表里 5 开头的）本月发生额转进 3103 本年利润；12 月 31 日把本年利润转进 310415 未分配利润'
    hp['A1'].font = Font(name='微软雅黑', sz=11, bold=True)
    hp['A2'] = ('A 列对齐科目表第 5 行起（损益科目的排序键）；B~F 是第 k 个损益科目（F＝上级科目直接记了账）；AM＝科目表期初；G~AL 是 32 个月各科目要结转的数'
                '（到本月的累计发生＋期初，四舍五入后减掉前几个月已结的；借-贷，不含结转凭证本身，只从【记账分录】第 4~16063 行取）；AN~BS 是累计条数；210 行起把每月有发生的科目挤到前面，记账分录的结转损益块按这里取数。')
    hp['A2'].font = f9
    for col, t in zip('ABCDEF', ['排序键', '科目表行', '科目编码', '科目全称', '分录条数', '非末级有账']):
        hp[f'{col}3'] = t
        hp[f'{col}3'].font = fb
    hp['B4'] = '最后有账月序：'
    hp['B4'].font = fb
    hp['C4'] = (f'=IF(COUNT({JE}!$A${JE0}:$A${JE1})=0,0,(YEAR(MAX({JE}!$A${JE0}:$A${JE1}))-YEAR(起始月份))*12'
                f'+MONTH(MAX({JE}!$A${JE0}:$A${JE1}))-MONTH(起始月份)+1)')
    for m in range(1, NM + 1):
        c, cc = mcol(m), cumcol(m)
        hp[f'{c}3'] = m
        hp[f'{c}4'] = f'=EDATE(起始月份,{m})-1'
        hp[f'{c}4'].number_format = 'yyyy-mm-dd'
        hp[f'{cc}3'] = m
        hp[f'{cc}4'] = 0
        for x in (f'{c}3', f'{cc}3'):
            hp[x].font = fb
    for r in range(5, 455):
        hp[f'A{r}'] = f'=IF(AND(科目表!$A{r}<>"",LEFT(科目表!$A{r},1)="5"),科目表!$N{r},"")'
    rr = f'{JE}!$R${JE0}:$R${JE1}'
    ss = f'{JE}!$S${JE0}:$S${JE1}'
    for k in range(1, NPL + 1):
        r = 4 + k
        hp[f'B{r}'] = f'=IFERROR(MATCH(SMALL($A$5:$A$454,{k}),$A$5:$A$454,0),"")'
        hp[f'C{r}'] = f'=IF($B{r}="","",INDEX(科目表!$A$5:$A$454,$B{r})&"")'
        hp[f'D{r}'] = f'=IF($B{r}="","",INDEX(科目表!$K$5:$K$454,$B{r}))'
        hp[f'AM{r}'] = f'=IF($C{r}="",0,N(SUMIF(科目表!$A$5:$A$454,$C{r},科目表!$M$5:$M$454)))'
        hp[f'E{r}'] = f'=IF($C{r}="",0,COUNTIFS({rr},">"&{k * 100},{rr},"<"&{k * 100 + 100})+IF($AM{r}<>0,1,0))'
        hp[f'F{r}'] = f'=IF($C{r}="",0,IF(AND($E{r}>0,INDEX(科目表!$L$5:$L$454,$B{r})="否"),1,0))'
        for m in range(1, NM + 1):
            c, cc = mcol(m), cumcol(m)
            # 结到本月为止的累计（含科目表期初）四舍五入后，减掉前几个月已经结掉的——尾差不会一个月一个月攒下来
            prev = '' if m == 1 else f'-SUM($G{r}:{mcol(m - 1)}{r})'
            hp[f'{c}{r}'] = (f'=IF(OR($E{r}=0,{c}$3>$C$4),0,ROUND(SUMIFS({ss},{rr},">"&{k * 100},{rr},"<="&({k * 100}+{c}$3))'
                             f'+$AM{r},2){prev})')
            hp[f'{c}{r}'].number_format = '#,##0.00;-#,##0.00;'
            hp[f'{cc}{r}'] = f'={cc}{r - 1}+IF({c}{r}<>0,1,0)'
    last = 4 + NPL
    labels = {R_TOT: '本月损益净额（借-贷）', R_CNT: '要结转的科目个数', R_YR: '年份', R_PL: '本年利润（借-贷）'}
    for r, t in labels.items():
        hp[f'F{r}'] = t
        hp[f'F{r}'].font = fb
    for m in range(1, NM + 1):
        c, cc = mcol(m), cumcol(m)
        hp[f'{c}{R_TOT}'] = f'=ROUND(SUM({c}5:{c}{last}),2)'
        hp[f'{c}{R_CNT}'] = f'={cc}{last}'
        hp[f'{c}{R_YR}'] = f'=YEAR({c}$4)'
        for j in range(1, NJ + 1):
            hp[f'{c}{R_K0 + j - 1}'] = f'=IFERROR(MATCH($F{R_K0 + j - 1},{cc}$5:{cc}${last},0),"")'
            hp[f'{c}{R_A0 + j - 1}'] = f'=IF({c}{R_K0 + j - 1}="",0,INDEX({c}$5:{c}${last},{c}{R_K0 + j - 1}))'
            hp[f'{c}{R_C0 + j - 1}'] = f'=IF({c}{R_K0 + j - 1}="","",INDEX($C$5:$C${last},{c}{R_K0 + j - 1}))'
        hp[f'{c}{R_PL}'] = f'=ROUND(SUM({c}{R_A0}:{c}{R_A0 + NJ - 1}),2)'
        for r in (R_TOT, R_PL) + tuple(range(R_A0, R_A0 + NJ)):
            hp[f'{c}{r}'].number_format = '#,##0.00;-#,##0.00;'
    for j in range(1, NJ + 1):
        for base, t in ((R_K0, '槽位'), (R_A0, '金额'), (R_C0, '科目')):
            hp[f'F{base + j - 1}'] = j
            hp[f'E{base + j - 1}'] = t if j == 1 else None
    for col, t in zip('BCDEF', ['年份', '年末日期', '本年利润手工发生(借-贷)', '各月结转(借-贷)', '年末余额(借-贷)']):
        hp[f'{col}{R_Y0 - 1}'] = t
        hp[f'{col}{R_Y0 - 1}'].font = fb
    for y in range(1, NY + 1):
        r = R_Y0 + y - 1
        hp[f'B{r}'] = f'=YEAR(起始月份)+{y - 1}'
        hp[f'C{r}'] = f'=DATE(B{r},12,31)'
        hp[f'C{r}'].number_format = 'yyyy-mm-dd'
        opening = 'N(SUMIF(科目表!$A$5:$A$454,"3103",科目表!$M$5:$M$454))+' if y == 1 else ''
        hp[f'D{r}'] = (f'=ROUND({opening}SUMIFS({JE}!$F${JE0}:$F${JE1},{JE}!$D${JE0}:$D${JE1},"3103",{JE}!$A${JE0}:$A${JE1},">="&DATE(B{r},1,1),'
                       f'{JE}!$A${JE0}:$A${JE1},"<="&C{r})-SUMIFS({JE}!$G${JE0}:$G${JE1},{JE}!$D${JE0}:$D${JE1},"3103",'
                       f'{JE}!$A${JE0}:$A${JE1},">="&DATE(B{r},1,1),{JE}!$A${JE0}:$A${JE1},"<="&C{r}),2)')
        hp[f'E{r}'] = f'=ROUND(SUMIF(${mcol(1)}${R_YR}:${mcol(NM)}${R_YR},B{r},${mcol(1)}${R_PL}:${mcol(NM)}${R_PL}),2)'
        hp[f'F{r}'] = f'=ROUND(D{r}+E{r},2)'
        for col in 'DEF':
            hp[f'{col}{r}'].number_format = '#,##0.00'
    for row in hp.iter_rows(min_row=3):
        for c in row:
            if c.value is not None and c.font.name != '微软雅黑':
                c.font = f9
    hp['AM3'] = '期初'
    hp['AM3'].font = fb
    hp.column_dimensions['D'].width = 28
    hp.freeze_panes = 'G5'
    log.append(f'新加隐藏表 {HP}（损益科目×32 个月的发生额、每月结转行、年末余额）')

    # 3) 记账分录：结转损益块（32 个月 × 61 行）＋ 年末结转块（3 年 × 2 行）
    tpl = {col: je[f'{col}16062'] for col in 'ABCDEFGHIJKLMNOPQ'}
    tpl2 = {col: je[f'{col}16063'] for col in 'ABCDEFGHIJKLMNOPQ'}

    def common(r, b, idx_off, idd):
        je[f'B{r}'] = f'=IF($P{r}=0,"",IFERROR(INDEX({IX}!$I$4:$I${IX_END},$L{r}-{idd}+{idx_off}),""))'
        je[f'E{r}'] = f'=IF($P{r}=0,"",IFERROR(INDEX(科目表!$K$5:$K$454,MATCH($D{r},科目表!$A$5:$A$454,0)),"（科目表里没有）"))'
        je[f'H{r}'] = '=""'
        je[f'I{r}'] = '=""'
        je[f'M{r}'] = f'=IF($P{r}=0,"",$Q${b}+SUM($P${b}:$P{r}))'
        je[f'N{r}'] = f'=IF($M{r}="","",$L{r}*1000+$M{r})'
        je[f'O{r}'] = (f'=IF($P{r}=0,"",IF(AND(明细账科目<>"",LEFT($D{r},LEN(明细账科目))=明细账科目,$A{r}>=明细账起,$A{r}<=明细账止),'
                       f'INT($A{r})*10^7+N($B{r})*1000+$M{r},""))')
        je[f'P{r}'] = f'=IF(AND($L{r}<>"",OR(ROUND($F{r},2)<>0,ROUND($G{r},2)<>0)),1,0)'
        je[f'Q{r}'] = 0 if r == b else None

    for m in range(1, NM + 1):
        c = mcol(m)
        b = CLOSE0 + (m - 1) * MBLK
        for j in range(1, MBLK + 1):
            r = b + j - 1
            for col, s in (tpl if j == 1 else tpl2).items():
                cp(s, je[f'{col}{r}'])
            je.row_dimensions[r].height = je.row_dimensions[16062].height
            je[f'A{r}'] = f'=IF($L{r}="","",{HP}!{c}$4)'
            je[f'C{r}'] = f'=IF($L{r}="","","结转"&MONTH($A{r})&"月损益")'
            if j <= NJ:
                je[f'D{r}'] = f'={HP}!{c}{R_C0 + j - 1}&""'
                je[f'F{r}'] = f'=MAX(-{HP}!{c}{R_A0 + j - 1},0)'
                je[f'G{r}'] = f'=MAX({HP}!{c}{R_A0 + j - 1},0)'
            else:                                   # 最后一行：本年利润
                je[f'D{r}'] = '3103'
                je[f'F{r}'] = f'=MAX({HP}!{c}${R_PL},0)'
                je[f'G{r}'] = f'=MAX(-{HP}!{c}${R_PL},0)'
            je[f'J{r}'] = '结转损益'
            je[f'K{r}'] = m
            je[f'L{r}'] = f'=IF(AND(ROUND($F{r},2)=0,ROUND($G{r},2)=0),"",{ID_M + m})'
            common(r, b, IXM0 - 4, ID_M)
    for y in range(1, NY + 1):
        b = YE0 + (y - 1) * 2
        hr = R_Y0 + y - 1
        for j, code in ((0, '3103'), (1, '310415')):
            r = b + j
            for col, s in (tpl if j == 0 else tpl2).items():
                cp(s, je[f'{col}{r}'])
            je.row_dimensions[r].height = je.row_dimensions[16062].height
            je[f'A{r}'] = f'=IF($L{r}="","",{HP}!$C${hr})'
            je[f'C{r}'] = f'=IF($L{r}="","","结转"&YEAR($A{r})&"年本年利润")'
            je[f'D{r}'] = code
            sgn = ('-', '') if j == 0 else ('', '-')
            je[f'F{r}'] = f'=MAX({sgn[0]}{HP}!$F${hr},0)'
            je[f'G{r}'] = f'=MAX({sgn[1]}{HP}!$F${hr},0)'
            je[f'J{r}'] = '年末结转'
            je[f'K{r}'] = y
            je[f'L{r}'] = f'=IF(AND(ROUND($F{r},2)=0,ROUND($G{r},2)=0),"",{ID_Y + y})'
            common(r, b, IXY0 - 4, ID_Y)
    assert je.auto_filter.ref == f'A3:J{JE1}', je.auto_filter.ref
    je.auto_filter.ref = f'A3:J{JE_END}'
    a2 = je['A2'].value
    s = '手工凭证、月末结转，每一笔'
    assert s in a2
    je['A2'].value = a2.replace(s, '手工凭证、月末结转（含每月结转损益、12 月 31 日结转本年利润），每一笔')
    log.append(f'记账分录 第 {CLOSE0}~{JE_END} 行：结转损益 32 个月 × {MBLK} 行、年末结转 3 年 × 2 行')

    # 4) _凭证索引 加 35 行
    ix = wb[IX]
    for r in range(IXM0, IX_END + 1):
        for col in 'ABCDEFGHIJKL':
            cp(ix[f'{col}{IX1}'], ix[f'{col}{r}'])
        if r < IXY0:
            m = r - IXM0 + 1
            b, e = CLOSE0 + (m - 1) * MBLK, CLOSE0 + m * MBLK - 1
            ix[f'A{r}'], ix[f'B{r}'] = ID_M + m, '结转损益'
            ix[f'D{r}'] = f'=IF($C{r}=0,"",{HP}!{mcol(m)}$4)'
            ix[f'J{r}'] = f'=IF($C{r}=0,"","结转"&MONTH($D{r})&"月损益")'
        else:
            y = r - IXY0 + 1
            b, e = YE0 + (y - 1) * 2, YE0 + (y - 1) * 2 + 1
            ix[f'A{r}'], ix[f'B{r}'] = ID_Y + y, '年末结转'
            ix[f'D{r}'] = f'=IF($C{r}=0,"",{HP}!$C${R_Y0 + y - 1})'
            ix[f'J{r}'] = f'=IF($C{r}=0,"","结转"&YEAR($D{r})&"年本年利润")'
        ix[f'C{r}'] = f'=IF(COUNTIF({JE}!$P${b}:$P${e},1)>0,1,0)'
        ix[f'E{r}'] = f'=IF($C{r}=0,"",INT($D{r})*10^7+$A{r})'
        ix[f'F{r}'] = f'=COUNTIF({JE}!$P${b}:$P${e},1)'
        ix[f'G{r}'] = f'=IF($C{r}=0,0,MAX(1,ROUNDUP($F{r}/8,0)))'
        ix[f'H{r}'] = f'=IF($C{r}=0,"",YEAR($D{r})*100+MONTH($D{r}))'
        ix[f'I{r}'] = f'=IF($C{r}=0,"",COUNTIFS({IX}!$H$4:$H${IX_END},$H{r},{IX}!$E$4:$E${IX_END},"<="&$E{r}))'
        ix[f'K{r}'] = '=""'
        ix[f'L{r}'] = f'=IF($C{r}=0,"",$H{r}*10000+$I{r})'
    log.append(f'_凭证索引 第 {IXM0}~{IX_END} 行：结转损益 32 张、年末结转 3 张')

    # 5) 范围扩到新的最后一行
    n = 0
    old_ix = f'{IX}!$I$4:$I${IX1}'
    for r in range(JE0, JE1 + 1):
        v = je[f'B{r}'].value
        assert old_ix in v, (r, v)
        je[f'B{r}'].value = v.replace(old_ix, f'{IX}!$I$4:$I${IX_END}')
        n += 1
    for r in range(4, IX1 + 1):
        v = ix[f'I{r}'].value
        w = v.replace(f'$H$4:$H${IX1}', f'$H$4:$H${IX_END}').replace(f'$E$4:$E${IX1}', f'$E$4:$E${IX_END}')
        assert w != v, (r, v)
        ix[f'I{r}'].value = w
        n += 1
    for nm, dn in list(wb.defined_names.items()):
        t = dn.attr_text
        if t.startswith(f'{JE}!') and t.endswith(f'${JE1}'):
            dn.attr_text = t[: -len(str(JE1))] + str(JE_END)
            n += 1
        elif t.startswith(f'{IX}!') and t.endswith(f'${IX1}'):
            dn.attr_text = t[: -len(str(IX1))] + str(IX_END)
            n += 1
    # 其他表里还有没有直接写死 6863 / 16063 的
    leftovers = []
    for ws in wb.worksheets:
        if ws.title in (JE, IX, HP):
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('=') and (f'${IX1}' in v or f'${JE1}' in v) and (IX in v or JE in v):
                    leftovers.append(f'{ws.title}!{c.coordinate}')
    assert not leftovers, leftovers[:10]
    log.append(f'记账分录、凭证索引的引用范围扩到 {JE_END} / {IX_END}（改了 {n} 处，含定义名称）')

    # 6) 科目余额表：利润表用的「不含结转损益」发生额（T~W，隐藏）
    kb = wb['科目余额表']
    hdrs = {'T': ('利润表用·本月(不含结转损益)', '借方'), 'U': (None, '贷方'), 'V': ('利润表用·本年(不含结转损益)', '借方'), 'W': (None, '贷方')}
    src = {'T': 'H', 'U': 'I', 'V': 'J', 'W': 'K'}
    for col, (h5, h6) in hdrs.items():
        for r in range(5, 458):
            assert kb[f'{col}{r}'].value is None
        cp(kb[f'{src[col]}5'], kb[f'{col}5'])
        cp(kb[f'{src[col]}6'], kb[f'{col}6'])
        kb[f'{col}5'].value, kb[f'{col}6'].value = h5, h6
        kb.column_dimensions[col].width = 14
        kb.column_dimensions[col].hidden = True
    bad = []
    for r in range(7, 457):
        for col, s in src.items():
            v = kb[f'{s}{r}'].value
            if not (isinstance(v, str) and v.startswith(f'=IF($Q{r}="","",SUMIFS(') and v.endswith(')))')):
                bad.append(f'{s}{r}: {v}')
                continue
            inner = v[len(f'=IF($Q{r}="","",'):-1]          # SUMIFS(...)
            kb[f'{col}{r}'] = f'=IF($Q{r}="","",IF(LEFT($A{r},1)="5",{inner[:-1]},分录来源,"<>结转损益"),{s}{r}))'
            cp(kb[f'{s}{r}'], kb[f'{col}{r}'])
    assert not bad, bad[:5]
    for nm, col in (('损益本月借', 'T'), ('损益本月贷', 'U'), ('损益本年借', 'V'), ('损益本年贷', 'W')):
        assert nm not in wb.defined_names
        wb.defined_names[nm] = DefinedName(nm, attr_text=f'科目余额表!${col}$7:${col}$456')
    log.append('科目余额表 T~W（隐藏）：损益科目本月 / 本年发生额，不含结转损益凭证')

    # 7) 利润表、经营报表改取「不含结转损益」的发生额
    rep = {'余额本月借': '损益本月借', '余额本月贷': '损益本月贷', '余额本年借': '损益本年借', '余额本年贷': '损益本年贷'}
    cnt = {}
    for sh in ('利润表', '经营报表'):
        ws = wb[sh]
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not (isinstance(v, str) and v.startswith('=')):
                    continue
                w = v
                for a, b2 in rep.items():
                    w = w.replace(a, b2)
                if sh == '经营报表':
                    w = re.sub(r'SUMIFS\((分录借方|分录贷方),分录科目,"5', r'SUMIFS(\1,分录来源,"<>结转损益",分录科目,"5', w)
                    # 别的写法的损益 SUMIFS（科目用格子引用的）要人看一下
                    for mm in re.finditer(r'SUMIFS\((?:分录借方|分录贷方),([^"]*?)分录科目,([^",)]+)', w):
                        raise AssertionError(f'{sh}!{c.coordinate} 有按格子取科目的 SUMIFS：{w[:200]}')
                if w != v:
                    c.value = w
                    cnt[sh] = cnt.get(sh, 0) + 1
    assert cnt.get('利润表', 0) >= 30 and cnt.get('经营报表', 0) >= 20, cnt
    log.append(f'利润表 {cnt["利润表"]} 格、经营报表 {cnt["经营报表"]} 格改取不含结转损益的发生额')

    # 8) 月末结转：每月汇总表加 J、K 两列
    mj = wb['月末结转']
    for r in range(5, 38):
        assert mj[f'J{r}'].value is None and mj[f'K{r}'].value is None
    for col, t in (('J', '结转损益\n本月利润\n(负＝亏损)'), ('K', '年末结转\n未分配利润')):
        cp(mj['I5'], mj[f'{col}5'])
        mj[f'{col}5'].value = t
    for m in range(1, NM + 1):
        r = 5 + m
        cp(mj[f'I{r}'], mj[f'J{r}'])
        cp(mj[f'I{r}'], mj[f'K{r}'])
        mj[f'J{r}'] = f'=-{HP}!{mcol(m)}${R_PL}'
        mj[f'K{r}'] = f'=IF(MONTH($B{r})=12,-IFERROR(INDEX({HP}!$F${R_Y0}:$F${R_Y0 + NY - 1},YEAR($B{r})-YEAR(起始月份)+1),0),"")'
    mj['A1'].value = '月末结转（折旧 · 调汇 · 生产成本结转原矿 · 结转销售成本 · 结转损益 · 年末结转本年利润）· 全自动'
    mj['A2'].value = (mj['A2'].value.rstrip('。') + '；⑥ 结转损益：每月最后一天把所有损益科目（收入、成本、税金、费用、营业外收支、所得税）'
                      '的本月发生额转进 3103 本年利润，损益科目月末余额为 0（J 列＝本月转进本年利润的利润，负数是亏损）；'
                      '⑦ 年末结转：12 月 31 日把本年利润全部转进 310415 利润分配-未分配利润（K 列）。'
                      '利润表、经营报表照常按本月 / 本年发生额出数（不含结转凭证），数不变。')
    mj.row_dimensions[2].height = 92
    log.append('月末结转 J/K 列：每月结转损益金额、12 月年末结转金额')


# =====================================================================
# ② 打印：A4 一张两页凭证
# =====================================================================
def step_print(wb, log):
    ws = wb['记账凭证']
    NB = 500
    assert ws['I1'].value == '=IF(1>本月总页数,"",1)' and ws['I15'].value == '=IF(2>本月总页数,"",2)'
    assert ws[f'I{14 * (NB - 1) + 1}'].value == f'=IF({NB}>本月总页数,"",{NB})'
    # 宽：原来 B~G 共 112 个字宽，A4 纵向印不下（贷方、√ 两列会跑到别的页）。收窄到 97 个字宽，缩放 75%：
    #     中文 Excel / WPS（宋体 11 每字宽 8 像素）印出来约 455 磅，LibreOffice 约 510 磅，A4 去掉页边能用 530 磅
    # 高：一页凭证 471 磅，两页＋中间 60 磅空＋末尾 6 磅＝1008 磅，×75%＝756 磅，A4 去掉页边能用 792 磅
    assert ws.column_dimensions['E'].min == 5 and ws.column_dimensions['E'].max == 6     # E:F 原来就是同一组列宽
    for col, w in (('B', 27), ('C', 12), ('D', 21), ('E', 16), ('G', 5)):
        ws.column_dimensions[col].width = w
    H = {1: 33.75, 2: 21.75, 3: 25.5, 12: 40.5, 13: 25.5}
    for i in range(4, 12):
        H[i] = 40.5
    fsz = {1: {'B': 18}, 2: {'B': 9, 'C': 11, 'E': 11, 'G': 9}, 3: {c: 11 for c in 'BCDEFG'},
           12: {c: 11 for c in 'BCEFG'}, 13: {c: 10 for c in 'BCDEF'}}
    for i in range(4, 12):
        fsz[i] = {'B': 10, 'C': 10.5, 'D': 10.5, 'E': 11, 'F': 11, 'G': 11}
    for blk in range(NB):
        base = 14 * blk
        for i in range(1, 14):
            ws.row_dimensions[base + i].height = H[i]
            for col, sz in fsz[i].items():
                c = ws[f'{col}{base + i}']
                f = copy(c.font)
                f.sz = sz
                c.font = f
            # 单位名、金额（上百亿的 13 位数）放不下时自动缩小字号，不会印成 ####
            if i == 2 or 4 <= i <= 12:
                for col in (('B',) if i == 2 else ('E', 'F')):
                    c = ws[f'{col}{base + i}']
                    al = copy(c.alignment)
                    al.shrink_to_fit, al.wrap_text = True, False
                    c.alignment = al
        ws.row_dimensions[base + 14].height = 60 if blk % 2 == 0 else 6
    ws.row_breaks = RowBreak()
    for p in range(1, NB // 2):
        ws.row_breaks.append(Break(id=28 * p))
    ps = ws.page_setup
    ps.paperSize = 9
    ps.orientation = 'portrait'
    ps.scale = 75
    ps.fitToWidth = ps.fitToHeight = None
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    pm = ws.page_margins
    pm.left, pm.right, pm.top, pm.bottom, pm.header, pm.footer = 0.6, 0.3, 0.4, 0.3, 0.2, 0.2
    ws.print_options.horizontalCentered = True
    s1 = 'I＝打印第几页'
    assert s1 in ws['S1'].value
    ws['S1'].value = ws['S1'].value.replace(s1, 'I＝凭证第几页（一张 A4 印两页：A4 第几张＝I÷2 进位，或看【凭证汇总】「在 A4 第几张」）')
    ws['S3'].value = ('月份跟【首页】的报表年度/月份走；一张 A4 纸上下印两页凭证（中间留宽，裁开后分别装订）。'
                      '打印时页码范围填 1 到【凭证汇总】右上角的「A4 张数」，后面的空页不用印。')
    ws['S4'].value = ('纸张已设好：A4 纵向、缩放 75%、左边留 1.5 厘米装订边。别改成「调整为一页」或「缩放到纸张」，不然两张凭证会挤在一起。'
                      '月末结转（折旧、调汇、成本、结转损益）和 12 月底的结转本年利润排在当月最后几号，跟着一起印。')
    cp(ws['S3'], ws['S4'])

    sm = wb['凭证汇总']
    a2 = sm['A2'].value
    s = '打印时页码范围填 1 到「总页数」。'
    assert s in a2
    sm['A2'].value = a2.replace(s, '一张 A4 纸上下印两页凭证：打印时页码范围填 1 到右上角「A4 张数」；只印某一张凭证，看它的「在 A4 第几张」。')
    assert sm['G3'].value == '总页数：' and sm['M3'].value is None and sm['N3'].value is None
    sm['G3'].value = '凭证页数：'
    cp(sm['G3'], sm['M3'])
    cp(sm['H3'], sm['N3'])
    sm['M3'].value = 'A4 张数：'
    sm['N3'].value = '=ROUNDUP(H3/2,0)'
    cp(sm['G3'], sm['O3'])
    sm['O3'].value = '="A4 张数："&N3'
    assert sm['O5'].value is None
    cp(sm['J5'], sm['O5'])
    sm['O5'].value = '在 A4\n第几张'
    for r in range(6, 456):
        cp(sm[f'J{r}'], sm[f'O{r}'])
        sm[f'O{r}'] = (f'=IF($J{r}="","",ROUNDUP($J{r}/2,0)&IF(ROUNDUP(($J{r}+$I{r}-1)/2,0)>ROUNDUP($J{r}/2,0),'
                       f'"～"&ROUNDUP(($J{r}+$I{r}-1)/2,0),""))')
        al = copy(sm[f'O{r}'].alignment)
        al.horizontal = 'center'
        sm[f'O{r}'].alignment = al
    sm.column_dimensions['O'].width = 13
    log.append('记账凭证：A4 纵向、缩放 75%，每 28 行一页（上下两页凭证，中间空 60 磅），列宽收窄、行高字号放大；凭证汇总加「A4 张数」「在 A4 第几张」')


# =====================================================================
# 数据校验、使用说明、首页
# =====================================================================
def step_check(wb, log):
    ck = wb['数据校验']
    assert ck['A42'].value == 38 and ck['A43'].value is None and ck['A44'].value is None
    rows = {
        43: (39, '结转损益：损益科目月末都结平了',
             '=IF(ROUND(SUMIFS(余额期末,余额编码,"5*",余额级次,1),2)=0,"√ 结平了（月末损益都转进了本年利润）",'
             '"✗ 损益科目月末还剩 "&FIXED(SUMIFS(余额期末,余额编码,"5*",余额级次,1),2)&"（看【科目余额表】5 开头哪个科目月末不是 0）")',
             '每月最后一天自动结转'),
        44: (40, '容量：结转损益每月最多 60 个科目',
             f'=IF(MAX({HP}!${mcol(1)}${R_CNT}:${mcol(NM)}${R_CNT})>{NJ},"✗ 有月份超过 {NJ} 个损益科目，多出来的没结转，找我加",'
             f'"√ 最多 "&MAX({HP}!${mcol(1)}${R_CNT}:${mcol(NM)}${R_CNT})&" 个")',
             f'每张结转损益凭证最多 {NJ} 行'),
    }
    for r, vals in rows.items():
        ck.row_dimensions[r].height = ck.row_dimensions[42].height
        for col, v in zip('ABCD', vals):
            cp(ck[f'{col}42'], ck[f'{col}{r}'])
            ck[f'{col}{r}'].value = v
    c29 = ck['C29'].value
    x = 'COUNTIFS(余额末级,"否",余额期末直接,"<>0")'
    assert c29.count(x) == 2
    ck['C29'].value = c29.replace(x, f'(COUNTIFS(余额末级,"否",余额期末直接,"<>0")+SUM({HP}!$F$5:$F${4 + NPL}))')
    old = ck['C3'].value
    assert old.count('C5:C42') == 3
    ck['C3'].value = old.replace('C5:C42', 'C5:C44')
    assert ck.auto_filter.ref == 'A4:D42'
    ck.auto_filter.ref = 'A4:D44'
    from openpyxl.formatting.formatting import ConditionalFormattingList
    cf_old = ck.conditional_formatting
    cf_new = ConditionalFormattingList()
    for rng in cf_old:
        ref = str(rng.sqref).replace('C5:C42', 'C5:C44')
        for rule in rng.rules:
            cf_new.add(ref, rule)
    ck.conditional_formatting = cf_new
    hm = wb['首页']
    v = hm['H11'].value
    assert v.count('C$5:$C$42') == 3
    hm['H11'].value = v.replace('C$5:$C$42', 'C$5:$C$44')
    log.append('数据校验加 2 项（损益结平、结转容量）')

    us = wb['使用说明']
    fixes = {
        'B5': ('一笔一张凭证；', '同一天、同一个账户的收款并成一张凭证（一借多贷），付款并成一张（多借一贷），苏姆现金、人民币现金、美元现金、各银行账户各自分开；'),
        'B6': ('月末结转（折旧、调汇、成本结转）全自动。', '月末结转（折旧、调汇、成本结转、结转损益到本年利润）和 12 月 31 日结转本年利润到未分配利润，全自动。'),
        'B15': ('「凭证合并号」写同一个号（同一个月）。', '「凭证合并号」写同一个号（同一个月）；不填就按天、按账户、按收付自动并。'),
        'B8': ('损益类科目不做结转（表结法）：利润表直接取发生额，',
               '损益类科目每月末自动结转（账结法）：月末最后一天「结转损益」把损益科目转进 3103 本年利润，12 月 31 日本年利润转进 310415 未分配利润；'
               '利润表取不含结转凭证的发生额，'),
        'B12': ('67 号矿井溢价款记采矿权；', '67 号矿井溢价款和退回的款记 122106 其他应收款-保证金（10/08 改）；'),
        'B25': ('【凭证汇总】看总页数 → 【记账凭证】打印，页码范围 1 到总页数（A5 横向，一页 8 条分录，超过的接下一页）。',
                '【凭证汇总】看「A4 张数」→ 【记账凭证】打印，页码范围 1 到 A4 张数（A4 纵向，一张纸上下两页凭证，中间留宽好裁开分别装订；'
                '一页凭证 8 条分录，超过的接下一页）。月末结转和 12 月底的年末结转凭证排在当月最后几号，一起印。'),
    }
    for ref, (a, b) in fixes.items():
        v = us[ref].value
        assert a in v, (ref, v)
        us[ref].value = v.replace(a, b)
    others = {
        ('科目余额表', 'A2'): ('损益类科目不结转（表结法），本年累计就是利润表的数；',
                            '损益类科目每月末结转进本年利润（账结法），所以损益科目月末余额是 0、发生额里含结转凭证；利润表取的是不含结转凭证的发生额；'),
        ('资产负债表', 'A2'): ('损益不结转（表结法），', '损益每月末结转进本年利润、12 月 31 日转进未分配利润，'),
        ('首页', 'F34'): ('打印记账凭证（A5 横向）', '打印记账凭证（A4 纵向，一张两页）'),
        ('首页', 'F35'): ('折旧、调汇、生产成本转原矿、结转销售成本', '折旧、调汇、生产成本转原矿、结转销售成本、结转损益、年末结转本年利润'),
        ('首页', 'B54'): ('生产成本转原矿、结转销售成本全自动，凭证日期是月末最后一天。',
                        '生产成本转原矿、结转销售成本、结转损益全自动，凭证日期是月末最后一天；12 月 31 日自动把本年利润转进未分配利润。'),
    }
    for (sh, ref), (a, b) in others.items():
        v = wb[sh][ref].value
        assert a in v, (sh, ref, v)
        wb[sh][ref].value = v.replace(a, b)
    from openpyxl.comments import Comment
    cs = wb[CASH]
    for r in SAVE_ROWS:
        old_c = cs[f'V{r}'].comment
        cm = Comment('10/08 改：67 号矿井溢价款（及退回的款）是保证金，记 122106 其他应收款-保证金（原来记 170101 采矿权）', '模板')
        cm.width, cm.height = (old_c.width, old_c.height) if old_c else (144, 79)
        cs[f'V{r}'].comment = cm
    log.append('使用说明、科目余额表、资产负债表、首页的说明文字和流水 6 行的批注改成新口径')


def no_col_overlap(wb):
    bad = []
    for ws in wb.worksheets:
        rng = sorted((d.min, d.max, k) for k, d in ws.column_dimensions.items() if d.min)
        for (a1, b1, k1), (a2, b2, k2) in zip(rng, rng[1:]):
            if a2 <= b1:
                bad.append(f'{ws.title}: {k1}({a1}-{b1}) 跟 {k2}({a2}-{b2}) 重叠')
    assert not bad, bad


def apply(wb):
    log = []
    step_deposit(wb, log)
    step_merge(wb, log)
    step_close(wb, log)
    step_print(wb, log)
    step_check(wb, log)
    no_col_overlap(wb)
    return log
