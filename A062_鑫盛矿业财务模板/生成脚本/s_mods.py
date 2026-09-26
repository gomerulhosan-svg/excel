# -*- coding: utf-8 -*-
"""业务录入模块：采购入库 / 领用出库 / 产量登记 / 销售结算 / 工资计提 / 固定资产 / 手工凭证。
现在都还没有数据（流水里的历史业务都是付款即记账），下个月开始出矿、销售就用得上。
每一行右边自动算出记账要用的科目；【记账分录】按固定位置把它们拼成凭证。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *
from s_base import code_of

N_DEPT_FA = 8          # 跟 s_mend.N_DEPT 一样：月末折旧按部门表前 8 个汇总

CODES = acc(AC_CODE)
b = lambda col, r0, r1: rng(SH_BASE, col, r0, r1)
MATN = b(M_NAME, MAT_R0, MAT_R1)
MAT = lambda col: b(col, MAT_R0, MAT_R1)


def _chk_fmt(ws, col, r0, r1):
    rg = f'{col}{r0}:{col}{r1}'
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="✗"'], fill=FILL_WARN,
                                                  font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="⚠"'], fill=fill('FFFFEB9C')))


def _rate(cur, d):
    return (f'IF(OR({cur}="",{cur}=本位币),1,IFERROR(LOOKUP(2,1/((汇率月份<={d})*(INDEX(记账汇率区,0,MATCH({cur},汇率币种头,0))<>""))'
            f',INDEX(记账汇率区,0,MATCH({cur},汇率币种头,0))),""))')


def month_no(d):
    """日期 → 月份序号（建账起始月份 = 1）"""
    return f'((YEAR({d})-YEAR(起始月份))*12+MONTH({d})-MONTH(起始月份)+1)'


def _seq(ws, r, blank):
    ws[f'A{r}'] = f'=IF({blank},"",ROW()-{MOD_R0 - 1})'


def _group_cols(ws, r, r1, valid, key_col, legs_col, pre_col, vid_col, key_expr, legs_expr, src):
    """凭证分组（隐藏列）：同一个分组键的几行并成一张凭证。
    key_expr：分组键（单据号；没填就用本行自己）；legs_expr：本行有几条非零分录。
    前面几行的分录条数 → 这一行的分录在凭证里从第几条接着排；凭证ID = 来源×100000 + 本组第一行的位置。"""
    ws[f'{key_col}{r}'] = f'=IF({valid}{r}=0,"",{key_expr})'
    ws[f'{legs_col}{r}'] = f'=IF({valid}{r}=0,0,{legs_expr})'
    ws[f'{pre_col}{r}'] = (f'=IF({valid}{r}=0,0,SUMIF(${key_col}${MOD_HDR}:{key_col}{r - 1},{key_col}{r},'
                           f'${legs_col}${MOD_HDR}:{legs_col}{r - 1}))')
    ws[f'{vid_col}{r}'] = (f'=IF({valid}{r}=0,"",{SRC[src] * 100000}+MATCH({key_col}{r},${key_col}${MOD_R0}:${key_col}${r1},0))')


def _group_style(ws, r0, r1, cols):
    for r in range(r0, r1 + 1):
        for c in cols:
            ws[f'{c}{r}'].font = F_NOTE
    hide(ws, cols)


def _chain(blank, pairs):
    """把 (条件, 提示) 串成嵌套 IF：第一个成立的条件给出提示，全不成立给空"""
    out = '""'
    for cond, msg in reversed(pairs):
        out = f'IF({cond},"{msg}",{out})' if not msg.startswith('=') else f'IF({cond},{msg[1:]},{out})'
    return f'=IF({blank},"",{out})'


def _same_month(dcol, r, head):
    """这一行跟本组第一行是不是同一个月"""
    return f'(YEAR({dcol}{r})*100+MONTH({dcol}{r})=YEAR({head})*100+MONTH({head}))'


def _head(col_date, col_no, r1, r):
    """同一个单据号的第一行的日期"""
    return f'INDEX(${col_date}${MOD_R0}:${col_date}${r1},MATCH({col_no}{r},${col_no}${MOD_R0}:${col_no}${r1},0))'


def _ok_group(col_date, col_no, r1, r):
    """没填单据号，或者跟同单据号第一行同一个月（✗ 的行不记账）"""
    return f'IF({col_no}{r}="",TRUE,IFERROR({_same_month(col_date, r, _head(col_date, col_no, r1, r))},FALSE))'


def _in_range(dcol, r):
    """日期是真日期、不早于建账起始月、在模板 32 个月以内"""
    return f'AND(ISNUMBER({dcol}{r}),{dcol}{r}>=起始月份,{month_no(f"{dcol}{r}")}<={N_MONTHS})'


# ───────────────────────── 采购入库 ─────────────────────────
def build_buy(wb, ctx):
    ws = wb.create_sheet(SH_BUY)
    widths(ws, {B_SEQ: 6, B_DATE: 11, B_NO: 11, B_SUP: 16, B_MAT: 14, B_SPEC: 9, B_UNIT: 6, B_QTY: 10, B_CUR: 6,
                B_UP: 12, B_AMT0: 14, B_RATE: 9, B_AMT: 15, B_FRT: 13, B_VAT: 13, B_COST: 15, B_TOT: 15, B_ACC: 9,
                B_PRJ: 10, B_NOTE: 16, B_CHK: 26, B_VALID: 5, B_GKEY: 6, B_LEGS: 5, B_GPRE: 5, B_VID: 8})
    title(ws, '采购入库（进口设备 · 炸药 · 油料 · 配件 · 五金 · 劳保 · 食材 …）', B_CHK, C_BUY,
          '💡 一张进货单一个物料一行：物料从下拉选（没有的先去【基础资料·物料档案】加），填数量、单价；外币（人民币/美元）选币种，自动按月度汇率折苏姆。'
          '运费、清关费、关税（付给海关、清关、运输公司的）填「运杂清关关税」，算进成本，挂「应付进口税费及运杂」；供应商发票上的增值税填「进项税」（海关交的进口增值税别填这里，在流水选「进口增值税（海关）」）。'
          '后来才到的运费、关税单子：再录一行同物料，只填「运杂清关关税」就行。同一张单子几个物料，「单据号」写一样的就并成一张凭证。'
          '自动凭证：借 存货/固定资产（按物料类别），借 进项税，贷 应付账款-供应商（货款＋进项税），贷 应付进口税费及运杂。'
          '付供应商在【现金流水】选「采购付款」（进口设备选「付设备款（进口）」）并写上同一个供应商；付关税运杂选「付进口税费及运杂」。')
    header(ws, MOD_HDR, [(B_SEQ, '序号'), (B_DATE, '日期'), (B_NO, '单据号'), (B_SUP, '供应商'), (B_MAT, '物料'),
                         (B_SPEC, '规格'), (B_UNIT, '单位'), (B_QTY, '数量'), (B_CUR, '币种'), (B_UP, '原币单价'),
                         (B_AMT0, '原币金额'), (B_RATE, '汇率'), (B_AMT, '折苏姆'), (B_FRT, '运杂清关关税\n(苏姆)'),
                         (B_VAT, '进项税\n(苏姆)'), (B_COST, '入账金额'), (B_TOT, '应付供应商\n(货款+进项税)'), (B_ACC, '入账科目'),
                         (B_PRJ, '项目'), (B_NOTE, '备注'), (B_CHK, '校验'), (B_VALID, '有效'), (B_GKEY, '分组'),
                         (B_LEGS, '条数'), (B_GPRE, '前面'), (B_VID, '凭证ID')], C_BUY, height=40)
    for r in range(MOD_R0, BUY_R1 + 1):
        blank = f'AND({B_DATE}{r}="",{B_SUP}{r}="",{B_MAT}{r}="",{B_QTY}{r}="",{B_UP}{r}="")'
        _seq(ws, r, blank)
        ws[f'{B_SPEC}{r}'] = f'=IF({B_MAT}{r}="","",IFERROR(INDEX({MAT(M_SPEC)},MATCH({B_MAT}{r},{MATN},0))&"",""))'
        ws[f'{B_UNIT}{r}'] = f'=IF({B_MAT}{r}="","",IFERROR(INDEX({MAT(M_UNIT)},MATCH({B_MAT}{r},{MATN},0))&"",""))'
        ws[f'{B_AMT0}{r}'] = f'=IF(OR(NOT(ISNUMBER({B_QTY}{r})),NOT(ISNUMBER({B_UP}{r}))),"",ROUND({B_QTY}{r}*{B_UP}{r},2))'
        ws[f'{B_RATE}{r}'] = f'=IF({B_AMT0}{r}="","",{_rate(f"{B_CUR}{r}", f"{B_DATE}{r}")})'
        ws[f'{B_AMT}{r}'] = f'=IF(OR({B_AMT0}{r}="",{B_RATE}{r}=""),"",ROUND({B_AMT0}{r}*{B_RATE}{r},2))'
        ws[f'{B_COST}{r}'] = f'=IF(AND({B_AMT}{r}="",{B_FRT}{r}=""),"",N({B_AMT}{r})+N({B_FRT}{r}))'
        ws[f'{B_TOT}{r}'] = f'=IF(AND({B_AMT}{r}="",{B_VAT}{r}=""),"",N({B_AMT}{r})+N({B_VAT}{r}))'
        ws[f'{B_ACC}{r}'] = f'=IF({B_MAT}{r}="","",IFERROR(INDEX({MAT(M_ACC)},MATCH({B_MAT}{r},{MATN},0))&"",""))'
        goods_ok = f'AND(ISNUMBER({B_QTY}{r}),ISNUMBER({B_UP}{r}),{B_RATE}{r}<>"")'
        addon_ok = (f'AND({B_QTY}{r}="",{B_UP}{r}="",OR(ISNUMBER({B_FRT}{r}),ISNUMBER({B_VAT}{r})))')
        extra_ok = f'AND(OR({B_FRT}{r}="",ISNUMBER({B_FRT}{r})),OR({B_VAT}{r}="",ISNUMBER({B_VAT}{r})))'
        ws[f'{B_VALID}{r}'] = (f'=IF(AND({_in_range(B_DATE, r)},N({B_COST}{r})+N({B_VAT}{r})<>0,{B_ACC}{r}<>"",'
                               f'OR({goods_ok},{addon_ok}),{extra_ok},'
                               f'ISNUMBER(MATCH({B_ACC}{r},{CODES},0)),{_ok_group(B_DATE, B_NO, BUY_R1, r)}),1,0)')
        _group_cols(ws, r, BUY_R1, B_VALID, B_GKEY, B_LEGS, B_GPRE, B_VID,
                    f'IF({B_NO}{r}<>"","D"&{B_NO}{r},"#"&ROW())',
                    f'(ROUND(N({B_COST}{r}),2)<>0)+(ROUND(N({B_VAT}{r}),2)<>0)+(ROUND(N({B_TOT}{r}),2)<>0)'
                    f'+(ROUND(N({B_FRT}{r}),2)<>0)', '采购')
        head = f'INDEX(${B_DATE}${MOD_R0}:${B_DATE}${BUY_R1},MATCH({B_NO}{r},${B_NO}${MOD_R0}:${B_NO}${BUY_R1},0))'
        ws[f'{B_CHK}{r}'] = _chain(blank, [
            (f'NOT(ISNUMBER({B_DATE}{r}))', '✗ 日期没填或不是真日期'), (f'{B_DATE}{r}<起始月份', '✗ 日期早于建账起始月份'),
            (f'{B_MAT}{r}=""', '✗ 没选物料'), (f'ISNA(MATCH({B_MAT}{r},{MATN},0))', '✗ 物料不在【基础资料·物料档案】里'),
            (f'{B_ACC}{r}=""', '✗ 物料没设类别（入账科目）'),
            (f'AND(OR({B_QTY}{r}="",{B_UP}{r}=""),NOT(AND({B_QTY}{r}="",{B_UP}{r}="",OR({B_FRT}{r}<>"",{B_VAT}{r}<>""))))',
             '✗ 没填数量或单价（只补运杂关税的行，数量单价都空着、只填运杂清关关税就行）'),
            (f'OR(AND({B_QTY}{r}<>"",NOT(ISNUMBER({B_QTY}{r}))),AND({B_UP}{r}<>"",NOT(ISNUMBER({B_UP}{r}))),'
             f'AND({B_FRT}{r}<>"",NOT(ISNUMBER({B_FRT}{r}))),AND({B_VAT}{r}<>"",NOT(ISNUMBER({B_VAT}{r}))))',
             '✗ 数量、单价、运杂、进项税要填数字（不要带单位）'),
            (f'{month_no(f"{B_DATE}{r}")}>{N_MONTHS}', '✗ 超出模板月份范围'),
            (f'AND({B_AMT0}{r}<>"",{B_RATE}{r}="")', '✗ 这个币种没有汇率（去【月度汇率】填）'),
            (f'ISNA(MATCH({B_ACC}{r},{CODES},0))', '✗ 入账科目不在科目表里'),
            (f'AND({B_NO}{r}<>"",NOT(IFERROR({_same_month(B_DATE, r, head)},TRUE)))', '✗ 同一个单据号的几行要在同一个月'),
            (f'{B_SUP}{r}=""', '⚠ 没填供应商（应付账款挂不到人）')])
    style_rows(ws, MOD_R0, BUY_R1, 'ABCDEFGHIJKLMNOPQRSTU',
               auto=(B_SEQ, B_SPEC, B_UNIT, B_AMT0, B_RATE, B_AMT, B_COST, B_TOT, B_ACC, B_CHK),
               fmts={B_DATE: DATE, B_QTY: QTY, B_UP: PRICE, B_AMT0: MONEY, B_RATE: RATE, B_AMT: MONEY, B_FRT: MONEY,
                     B_VAT: MONEY, B_COST: MONEY, B_TOT: MONEY},
               aligns={B_NOTE: AL, B_CHK: AL, **{c: AR_ for c in (B_QTY, B_UP, B_AMT0, B_RATE, B_AMT, B_FRT, B_VAT,
                                                                  B_COST, B_TOT)}})
    _group_style(ws, MOD_R0, BUY_R1, [B_VALID, B_GKEY, B_LEGS, B_GPRE, B_VID])
    _chk_fmt(ws, B_CHK, MOD_R0, BUY_R1)
    add_date_dv(ws, f'{B_DATE}{MOD_R0}:{B_DATE}{BUY_R1}')
    add_list_dv(ws, f'{B_SUP}{MOD_R0}:{B_SUP}{BUY_R1}', '=往来列表', stop=False)
    add_list_dv(ws, f'{B_MAT}{MOD_R0}:{B_MAT}{BUY_R1}', '=物料列表')
    add_list_dv(ws, f'{B_CUR}{MOD_R0}:{B_CUR}{BUY_R1}', '=汇率币种头', '空着＝苏姆')
    add_list_dv(ws, f'{B_PRJ}{MOD_R0}:{B_PRJ}{BUY_R1}', '=项目列表', stop=False)
    ws.auto_filter.ref = f'A{MOD_HDR}:{B_CHK}{BUY_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 领用出库 ─────────────────────────
def build_issue(wb, ctx):
    ws = wb.create_sheet(SH_ISS)
    widths(ws, {I_SEQ: 6, I_DATE: 11, I_NO: 11, I_MAT: 14, I_SPEC: 9, I_UNIT: 6, I_QTY: 10, I_USE: 12, I_PRICE: 13,
                I_AMT: 15, I_DR: 10, I_CR: 9, I_PRJ: 10, I_NOTE: 18, I_CHK: 26, I_VALID: 5, I_GKEY: 6, I_LEGS: 5,
                I_GPRE: 5, I_VID: 8})
    title(ws, '领用出库（炸药、油料、配件、五金、劳保…从库里领出来用）', I_CHK, C_ISS,
          '💡 领一次记一行：选物料、填数量、选用途。单价自动按「全月一次加权平均」（本月期初＋本月进货）算，金额自动。'
          '用途决定记到哪：采矿生产→生产成本-直接材料（炸药/油料/配件分开），设备维修→生产成本-机械使用费，基建掘进→在建工程-井巷工程，'
          '行政、食堂→管理费用（在【记账规则·领用用途】里能改）。自动凭证：借 用途科目，贷 原材料/周转材料。'
          '「领料单号」写一样的几行并成一张凭证（比如整个月写同一个号「2026-09」就一个月一张）。')
    header(ws, MOD_HDR, [(I_SEQ, '序号'), (I_DATE, '日期'), (I_NO, '领料单号'), (I_MAT, '物料'), (I_SPEC, '规格'),
                         (I_UNIT, '单位'), (I_QTY, '数量'), (I_USE, '用途'), (I_PRICE, '单价\n(月加权平均)'),
                         (I_AMT, '金额'), (I_DR, '计入科目'), (I_CR, '存货科目'), (I_PRJ, '项目'), (I_NOTE, '领用人/备注'),
                         (I_CHK, '校验'), (I_VALID, '有效'), (I_GKEY, '分组'), (I_LEGS, '条数'), (I_GPRE, '前面'),
                         (I_VID, '凭证ID')], C_ISS, height=40)
    us = lambda col: rng(SH_RULE, col, US_R0, US_R1)
    mc = lambda col: rng(SH_RULE, col, MC_R0, MC_R1)
    from s_mend import inv_unit_cost
    for r in range(MOD_R0, ISS_R1 + 1):
        blank = f'AND({I_DATE}{r}="",{I_MAT}{r}="",{I_QTY}{r}="")'
        _seq(ws, r, blank)
        ws[f'{I_SPEC}{r}'] = f'=IF({I_MAT}{r}="","",IFERROR(INDEX({MAT(M_SPEC)},MATCH({I_MAT}{r},{MATN},0))&"",""))'
        ws[f'{I_UNIT}{r}'] = f'=IF({I_MAT}{r}="","",IFERROR(INDEX({MAT(M_UNIT)},MATCH({I_MAT}{r},{MATN},0))&"",""))'
        ws[f'{I_PRICE}{r}'] = (f'=IF(OR({I_MAT}{r}="",NOT(ISNUMBER({I_DATE}{r}))),"",'
                               f'IFERROR({inv_unit_cost(f"MATCH({I_MAT}{r},{MATN},0)", month_no(f"{I_DATE}{r}"))},""))')
        ws[f'{I_AMT}{r}'] = f'=IF(OR(NOT(ISNUMBER({I_QTY}{r})),{I_PRICE}{r}=""),"",ROUND({I_QTY}{r}*{I_PRICE}{r},2))'
        cat = f'INDEX({MAT(M_CAT)},MATCH({I_MAT}{r},{MATN},0))'
        u = f'INDEX({us(US_ACC)},MATCH({I_USE}{r},{us(US_NAME)},0))&""'
        prod = f'INDEX({mc(MC_PROD)},MATCH({cat},{mc(MC_NAME)},0))&""'
        ws[f'{I_DR}{r}'] = (f'=IF(OR({I_USE}{r}="",{I_MAT}{r}=""),"",IFERROR(IF({u}="按物料类别",{prod},'
                            f'{code_of(u)}),""))')
        ws[f'{I_CR}{r}'] = f'=IF({I_MAT}{r}="","",IFERROR(INDEX({MAT(M_ACC)},MATCH({I_MAT}{r},{MATN},0))&"",""))'
        stock = f'IFERROR(INDEX({MAT(M_STOCK)},MATCH({I_MAT}{r},{MATN},0)),"")'
        ws[f'{I_VALID}{r}'] = (f'=IF(AND({_in_range(I_DATE, r)},N({I_AMT}{r})<>0,{I_DR}{r}<>"",{I_CR}{r}<>"",'
                               f'ISNUMBER(MATCH({I_DR}{r},{CODES},0)),ISNUMBER(MATCH({I_CR}{r},{CODES},0)),{stock}="是",'
                               f'{_ok_group(I_DATE, I_NO, ISS_R1, r)}),1,0)')
        _group_cols(ws, r, ISS_R1, I_VALID, I_GKEY, I_LEGS, I_GPRE, I_VID,
                    f'IF({I_NO}{r}<>"","D"&{I_NO}{r},"#"&ROW())', '2', '领用')
        head = f'INDEX(${I_DATE}${MOD_R0}:${I_DATE}${ISS_R1},MATCH({I_NO}{r},${I_NO}${MOD_R0}:${I_NO}${ISS_R1},0))'
        ws[f'{I_CHK}{r}'] = _chain(blank, [
            (f'NOT(ISNUMBER({I_DATE}{r}))', '✗ 日期没填或不是真日期'), (f'{I_DATE}{r}<起始月份', '✗ 日期早于建账起始月份'),
            (f'{month_no(f"{I_DATE}{r}")}>{N_MONTHS}', '✗ 超出模板月份范围'),
            (f'{I_MAT}{r}=""', '✗ 没选物料'), (f'ISNA(MATCH({I_MAT}{r},{MATN},0))', '✗ 物料不在物料档案里'),
            (f'{stock}<>"是"', '✗ 这个物料不是库存管理的（买来就直接记账了），不用领用'),
            (f'{I_QTY}{r}=""', '✗ 没填数量'), (f'NOT(ISNUMBER({I_QTY}{r}))', '✗ 数量要填数字（不要带单位）'),
            (f'{I_USE}{r}=""', '✗ 没选用途'),
            (f'{I_DR}{r}=""', '✗ 用途在【记账规则】里没设科目'),
            (f'OR(ISNA(MATCH({I_DR}{r},{CODES},0)),ISNA(MATCH({I_CR}{r},{CODES},0)))', '✗ 计入科目或存货科目不在科目表里'),
            (f'N({I_PRICE}{r})=0', '⚠ 库里没有这个物料的进货（单价为 0），先录采购入库'),
            (f'AND({I_NO}{r}<>"",NOT(IFERROR({_same_month(I_DATE, r, head)},TRUE)))', '✗ 同一个领料单号的几行要在同一个月')])
    style_rows(ws, MOD_R0, ISS_R1, 'ABCDEFGHIJKLMNO', auto=(I_SEQ, I_SPEC, I_UNIT, I_PRICE, I_AMT, I_DR, I_CR, I_CHK),
               fmts={I_DATE: DATE, I_QTY: QTY, I_PRICE: PRICE, I_AMT: MONEY},
               aligns={I_NOTE: AL, I_CHK: AL, I_QTY: AR_, I_PRICE: AR_, I_AMT: AR_})
    _group_style(ws, MOD_R0, ISS_R1, [I_VALID, I_GKEY, I_LEGS, I_GPRE, I_VID])
    _chk_fmt(ws, I_CHK, MOD_R0, ISS_R1)
    add_date_dv(ws, f'{I_DATE}{MOD_R0}:{I_DATE}{ISS_R1}')
    add_list_dv(ws, f'{I_MAT}{MOD_R0}:{I_MAT}{ISS_R1}', '=物料列表')
    add_list_dv(ws, f'{I_USE}{MOD_R0}:{I_USE}{ISS_R1}', '=用途列表')
    add_list_dv(ws, f'{I_PRJ}{MOD_R0}:{I_PRJ}{ISS_R1}', '=项目列表', stop=False)
    ws.auto_filter.ref = f'A{MOD_HDR}:{I_CHK}{ISS_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 产量登记 ─────────────────────────
def build_prod(wb, ctx):
    ws = wb.create_sheet(SH_PROD)
    widths(ws, {PD_SEQ: 6, PD_DATE: 11, PD_FACE: 16, PD_TON: 14, PD_DEV: 12, PD_GRADE: 10, PD_DIL: 10, PD_NOTE: 24,
                PD_CHK: 24})
    title(ws, '产量登记（出矿量 · 掘进量 · 品位）', PD_CHK, C_PROD,
          '💡 按天或按班记都行，一个作业面一行：出矿量（吨，以过磅/验收为准）、掘进进尺（米）、品位、贫化率。'
          '每月月底自动把当月生产成本（炸药、油料、配件、人工、折旧…）结转成原矿库存成本（有出矿的月份才结转），'
          '【原矿产销存】算出每吨成本，卖出去时自动结转销售成本。【经营报表】的出矿量、掘进量、吨矿成本也从这里取。')
    header(ws, MOD_HDR, [(PD_SEQ, '序号'), (PD_DATE, '日期'), (PD_FACE, '作业面/采区'), (PD_TON, '出矿量(吨)'),
                         (PD_DEV, '掘进量(米)'), (PD_GRADE, '品位(%)'), (PD_DIL, '贫化率(%)'), (PD_NOTE, '备注'),
                         (PD_CHK, '校验')], C_PROD, height=36)
    for r in range(MOD_R0, PROD_R1 + 1):
        blank = f'AND({PD_DATE}{r}="",{PD_TON}{r}="",{PD_DEV}{r}="")'
        _seq(ws, r, blank)
        ws[f'{PD_CHK}{r}'] = (f'=IF({blank},"",IF(NOT(ISNUMBER({PD_DATE}{r})),"✗ 日期没填或不是真日期",'
                              f'IF({PD_DATE}{r}<起始月份,"✗ 日期早于建账起始月份",'
                              f'IF(AND({PD_TON}{r}<>"",NOT(ISNUMBER({PD_TON}{r}))),"✗ 出矿量不是数字",""))))')
    style_rows(ws, MOD_R0, PROD_R1, 'ABCDEFGHI', auto=(PD_SEQ, PD_CHK),
               fmts={PD_DATE: DATE, PD_TON: '#,##0.00', PD_DEV: '#,##0.0', PD_GRADE: '0.00', PD_DIL: '0.00'},
               aligns={PD_NOTE: AL, PD_CHK: AL, PD_TON: AR_, PD_DEV: AR_})
    _chk_fmt(ws, PD_CHK, MOD_R0, PROD_R1)
    add_date_dv(ws, f'{PD_DATE}{MOD_R0}:{PD_DATE}{PROD_R1}')
    ws.auto_filter.ref = f'A{MOD_HDR}:{PD_CHK}{PROD_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 销售结算 ─────────────────────────
def build_sale(wb, ctx):
    ws = wb.create_sheet(SH_SALE)
    widths(ws, {S_SEQ: 6, S_DATE: 11, S_NO: 13, S_CUS: 16, S_ITEM: 8, S_TON: 12, S_PRICE: 13, S_AMT: 16, S_RATE: 7,
                S_NET: 16, S_TAX: 15, S_TOT: 16, S_UC: 13, S_COGS: 16, S_PRJ: 10, S_NOTE: 16, S_CHK: 24, S_VALID: 5,
                S_GKEY: 6, S_LEGS: 5, S_GPRE: 5, S_VID: 8})
    title(ws, '销售结算（原矿按吨卖给甲方）', S_CHK, C_SALE,
          '💡 一张结算单/磅单一行：客户（甲方）、吨数、单价（苏姆/吨）。单价是不是含税看【基础资料】参数（现在是「含税」），'
          '税率默认 12%（这一格填了就按填的）。自动凭证：借 应收账款-原矿销售（价税合计），贷 主营业务收入-原矿销售、应交税费-销项税额；'
          '同时按【原矿产销存】每吨成本结转销售成本（借 主营业务成本，贷 库存商品-原矿）。甲方付款在【现金流水】选「销售回款」、往来单位写甲方。')
    header(ws, MOD_HDR, [(S_SEQ, '序号'), (S_DATE, '日期'), (S_NO, '结算单/磅单号'), (S_CUS, '客户（甲方）'),
                         (S_ITEM, '品名'), (S_TON, '数量(吨)'), (S_PRICE, '单价\n(苏姆/吨)'), (S_AMT, '金额'),
                         (S_RATE, '税率'), (S_NET, '不含税收入'), (S_TAX, '销项税'), (S_TOT, '价税合计\n(应收)'),
                         (S_UC, '每吨成本\n(自动)'), (S_COGS, '销售成本\n(自动)'), (S_PRJ, '项目'), (S_NOTE, '备注'),
                         (S_CHK, '校验'), (S_VALID, '有效'), (S_GKEY, '分组'), (S_LEGS, '条数'), (S_GPRE, '前面'),
                         (S_VID, '凭证ID')], C_SALE, height=40)
    from s_mend import ore_unit_cost
    for r in range(MOD_R0, SALE_R1 + 1):
        blank = f'AND({S_DATE}{r}="",{S_CUS}{r}="",{S_TON}{r}="",{S_PRICE}{r}="")'
        _seq(ws, r, blank)
        ws[f'{S_AMT}{r}'] = f'=IF(OR(NOT(ISNUMBER({S_TON}{r})),NOT(ISNUMBER({S_PRICE}{r}))),"",ROUND({S_TON}{r}*{S_PRICE}{r},2))'
        rt = f'IF({S_RATE}{r}="",增值税率,{S_RATE}{r})'
        ws[f'{S_NET}{r}'] = f'=IF({S_AMT}{r}="","",IF(单价含税="是",ROUND({S_AMT}{r}/(1+{rt}),2),{S_AMT}{r}))'
        ws[f'{S_TAX}{r}'] = f'=IF({S_AMT}{r}="","",IF(单价含税="是",{S_AMT}{r}-{S_NET}{r},ROUND({S_AMT}{r}*{rt},2)))'
        ws[f'{S_TOT}{r}'] = f'=IF({S_AMT}{r}="","",{S_NET}{r}+{S_TAX}{r})'
        ws[f'{S_UC}{r}'] = f'=IF(NOT(ISNUMBER({S_DATE}{r})),"",IFERROR({ore_unit_cost(month_no(f"{S_DATE}{r}"))},0))'
        ws[f'{S_COGS}{r}'] = f'=IF(OR({S_TON}{r}="",{S_UC}{r}=""),"",ROUND({S_TON}{r}*{S_UC}{r},2))'
        ws[f'{S_VALID}{r}'] = (f'=IF(AND({_in_range(S_DATE, r)},N({S_TOT}{r})<>0,'
                               f'{_ok_group(S_DATE, S_NO, SALE_R1, r)}),1,0)')
        _group_cols(ws, r, SALE_R1, S_VALID, S_GKEY, S_LEGS, S_GPRE, S_VID,
                    f'IF({S_NO}{r}<>"","D"&{S_NO}{r},"#"&ROW())',
                    f'(ROUND(N({S_NET}{r}),2)<>0)+(ROUND(N({S_TAX}{r}),2)<>0)+(ROUND(N({S_TOT}{r}),2)<>0)', '销售')
        head = f'INDEX(${S_DATE}${MOD_R0}:${S_DATE}${SALE_R1},MATCH({S_NO}{r},${S_NO}${MOD_R0}:${S_NO}${SALE_R1},0))'
        ws[f'{S_CHK}{r}'] = _chain(blank, [
            (f'NOT(ISNUMBER({S_DATE}{r}))', '✗ 日期没填或不是真日期'), (f'{S_DATE}{r}<起始月份', '✗ 日期早于建账起始月份'),
            (f'{month_no(f"{S_DATE}{r}")}>{N_MONTHS}', '✗ 超出模板月份范围'),
            (f'OR({S_TON}{r}="",{S_PRICE}{r}="")', '✗ 没填吨数或单价（没定价的先别录，定了价再录）'),
            (f'OR(NOT(ISNUMBER({S_TON}{r})),NOT(ISNUMBER({S_PRICE}{r})),AND({S_RATE}{r}<>"",NOT(ISNUMBER({S_RATE}{r}))))',
             '✗ 吨数、单价、税率要填数字（不要带单位）'),
            (f'AND({S_NO}{r}<>"",NOT(IFERROR({_same_month(S_DATE, r, head)},TRUE)))', '✗ 同一个结算单号的几行要在同一个月'),
            (f'{S_CUS}{r}=""', '⚠ 没填客户（应收账款挂不到甲方）'),
            (f'N({S_UC}{r})=0', '⚠ 这个月原矿没有成本（没登记产量或没结转），销售成本按 0')])
    style_rows(ws, MOD_R0, SALE_R1, 'ABCDEFGHIJKLMNOPQ', auto=(S_SEQ, S_AMT, S_NET, S_TAX, S_TOT, S_UC, S_COGS, S_CHK),
               fmts={S_DATE: DATE, S_TON: '#,##0.00', S_PRICE: MONEY, S_AMT: MONEY, S_RATE: '0%', S_NET: MONEY,
                     S_TAX: MONEY, S_TOT: MONEY, S_UC: MONEY, S_COGS: MONEY},
               aligns={S_NOTE: AL, S_CHK: AL, **{c: AR_ for c in (S_TON, S_PRICE, S_AMT, S_NET, S_TAX, S_TOT, S_UC, S_COGS)}})
    _group_style(ws, MOD_R0, SALE_R1, [S_VALID, S_GKEY, S_LEGS, S_GPRE, S_VID])
    _chk_fmt(ws, S_CHK, MOD_R0, SALE_R1)
    add_date_dv(ws, f'{S_DATE}{MOD_R0}:{S_DATE}{SALE_R1}')
    add_list_dv(ws, f'{S_CUS}{MOD_R0}:{S_CUS}{SALE_R1}', '=往来列表', stop=False)
    add_list_dv(ws, f'{S_PRJ}{MOD_R0}:{S_PRJ}{SALE_R1}', '=项目列表', stop=False)
    ws.auto_filter.ref = f'A{MOD_HDR}:{S_CHK}{SALE_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 工资计提 ─────────────────────────
def build_pay(wb, ctx):
    ws = wb.create_sheet(SH_PAY)
    widths(ws, {W_SEQ: 6, W_DATE: 11, W_DEPT: 12, W_TYPE: 9, W_HEAD: 7, W_AMT: 16, W_DR: 10, W_CR: 9, W_NOTE: 24,
                W_CHK: 24, W_VALID: 5, W_GKEY: 7, W_LEGS: 5, W_GPRE: 5, W_VID: 8})
    title(ws, '工资计提（每月按部门、中籍/外籍各一行）', W_CHK, C_PAY,
          '💡 每月底记一次：部门、人员类别、人数、应发工资（苏姆）。自动凭证：借 部门对应的成本费用科目（采矿生产→生产成本-直接人工、'
          '基建→在建工程、行政→管理费用），贷 应付职工薪酬。发工资时在【现金流水】选「应付职工薪酬」冲掉。'
          '同一个月的几行自动并成一张计提凭证。（业务模块启用日以前的工资没计提，发的时候直接记了管理费用，跟你 8 月报表一样。）')
    header(ws, MOD_HDR, [(W_SEQ, '序号'), (W_DATE, '日期'), (W_DEPT, '部门'), (W_TYPE, '人员类别'), (W_HEAD, '人数'),
                         (W_AMT, '应发工资(苏姆)'), (W_DR, '计入科目'), (W_CR, '应付科目'), (W_NOTE, '备注'),
                         (W_CHK, '校验'), (W_VALID, '有效'), (W_GKEY, '月份'), (W_LEGS, '条数'), (W_GPRE, '前面'),
                         (W_VID, '凭证ID')], C_PAY, height=36)
    dp = lambda col: rng(SH_RULE, col, DP_R0, DP_R1)
    pt = lambda col: rng(SH_RULE, col, PT_R0, PT_R1)
    for r in range(MOD_R0, PAY_R1 + 1):
        blank = f'AND({W_DATE}{r}="",{W_DEPT}{r}="",{W_AMT}{r}="")'
        _seq(ws, r, blank)
        ws[f'{W_DR}{r}'] = f'=IF({W_DEPT}{r}="","",IFERROR({code_of(f"INDEX({dp(DP_WAGE)},MATCH({W_DEPT}{r},{dp(DP_NAME)},0))")},""))'
        ws[f'{W_CR}{r}'] = f'=IF({W_TYPE}{r}="","221101",IFERROR({code_of(f"INDEX({pt(PT_ACC)},MATCH({W_TYPE}{r},{pt(PT_NAME)},0))")},""))'
        ws[f'{W_VALID}{r}'] = (f'=IF(AND({_in_range(W_DATE, r)},ISNUMBER({W_AMT}{r}),N({W_AMT}{r})<>0,'
                               f'ISNUMBER(MATCH({W_DR}{r},{CODES},0)),ISNUMBER(MATCH({W_CR}{r},{CODES},0))),1,0)')
        _group_cols(ws, r, PAY_R1, W_VALID, W_GKEY, W_LEGS, W_GPRE, W_VID,
                    f'YEAR({W_DATE}{r})*100+MONTH({W_DATE}{r})', '2', '工资')
        ws[f'{W_CHK}{r}'] = _chain(blank, [
            (f'NOT(ISNUMBER({W_DATE}{r}))', '✗ 日期没填或不是真日期'), (f'{W_DATE}{r}<起始月份', '✗ 日期早于建账起始月份'),
            (f'{W_DEPT}{r}=""', '✗ 没选部门'), (f'{W_DR}{r}=""', '✗ 这个部门在【记账规则】里没设工资科目'),
            (f'OR(ISNA(MATCH({W_DR}{r},{CODES},0)),ISNA(MATCH({W_CR}{r},{CODES},0)))', '✗ 工资科目不在科目表里'),
            (f'{W_AMT}{r}=""', '✗ 没填工资金额'), (f'NOT(ISNUMBER({W_AMT}{r}))', '✗ 工资金额要填数字'),
            (f'{month_no(f"{W_DATE}{r}")}>{N_MONTHS}', '✗ 超出模板月份范围')])
    style_rows(ws, MOD_R0, PAY_R1, 'ABCDEFGHIJ', auto=(W_SEQ, W_DR, W_CR, W_CHK),
               fmts={W_DATE: DATE, W_HEAD: '0', W_AMT: MONEY}, aligns={W_NOTE: AL, W_CHK: AL, W_AMT: AR_})
    _group_style(ws, MOD_R0, PAY_R1, [W_VALID, W_GKEY, W_LEGS, W_GPRE, W_VID])
    _chk_fmt(ws, W_CHK, MOD_R0, PAY_R1)
    add_date_dv(ws, f'{W_DATE}{MOD_R0}:{W_DATE}{PAY_R1}')
    add_list_dv(ws, f'{W_DEPT}{MOD_R0}:{W_DEPT}{PAY_R1}', '=部门列表')
    add_list_dv(ws, f'{W_TYPE}{MOD_R0}:{W_TYPE}{PAY_R1}', '=人员类别列表')
    ws.auto_filter.ref = f'A{MOD_HDR}:{W_CHK}{PAY_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 固定资产 ─────────────────────────
def build_fa(wb, ctx):
    ws = wb.create_sheet(SH_FA)
    widths(ws, {F_NO: 8, F_NAME: 20, F_CAT: 12, F_ACC: 9, F_COST: 16, F_DATE: 11, F_YEARS: 7, F_SALV: 7, F_DEPT: 10,
                F_DACC: 10, F_MDEP: 14, F_START: 10, F_MONTHS: 8, F_ACCUM: 16, F_NBV: 16, F_REMK: 16, F_CHK: 24,
                F_YEFF: 5, F_SEFF: 5, F_END: 9, F_LASTAMT: 12, F_OK: 5})
    title(ws, '固定资产卡片（矿车 · 四不像 · 铲车 · 挖掘机 · 发电机组 …）· 自动提折旧', F_CHK, C_FA,
          '💡 一台设备一行：原值（苏姆，含运费清关关税）、入账日期、类别、使用部门。年限、残值率不填就按【记账规则·固定资产类别】的默认。'
          '入账的下个月开始按直线法每月自动提折旧（【月末结转】出折旧凭证：采矿生产→生产成本-折旧费，基建→在建工程，行政→管理费用-折旧费）。'
          '注意：卡片只管折旧，设备买进来的那张凭证是在【采购入库】或【现金流水】记的，这里不会重复记。')
    header(ws, MOD_HDR, [(F_NO, '编号'), (F_NAME, '名称/型号'), (F_CAT, '类别'), (F_ACC, '科目'), (F_COST, '原值(苏姆)'),
                         (F_DATE, '入账日期'), (F_YEARS, '年限\n(选填)'), (F_SALV, '残值率\n(选填)'), (F_DEPT, '使用部门'),
                         (F_DACC, '折旧计入'), (F_MDEP, '月折旧额'), (F_START, '开始折旧'),
                         (F_MONTHS, '截至报表\n月已提'), (F_ACCUM, '累计折旧'), (F_NBV, '净值'), (F_REMK, '备注'),
                         (F_CHK, '校验'), (F_YEFF, '年限'), (F_SEFF, '残值'), (F_END, '最后折旧月'),
                         (F_LASTAMT, '末月折旧'), (F_OK, '没错')], C_FA, height=40)
    fc = lambda col: rng(SH_RULE, col, FC_R0, FC_R1)
    dp = lambda col: rng(SH_RULE, col, DP_R0, DP_R1)
    rep_end = 'DATE(报表年度,报表月份+1,0)'
    for r in range(MOD_R0, FA_R1 + 1):
        blank = f'AND({F_NAME}{r}="",{F_COST}{r}="")'
        ws[f'{F_ACC}{r}'] = f'=IF({F_CAT}{r}="","",IFERROR({code_of(f"INDEX({fc(FC_ACC)},MATCH({F_CAT}{r},{fc(FC_NAME)},0))")},""))'
        ws[f'{F_YEFF}{r}'] = f'=IF({F_YEARS}{r}<>"",{F_YEARS}{r},IFERROR(INDEX({fc(FC_YEARS)},MATCH({F_CAT}{r},{fc(FC_NAME)},0)),""))'
        ws[f'{F_SEFF}{r}'] = f'=IF({F_SALV}{r}<>"",{F_SALV}{r},IFERROR(INDEX({fc(FC_SALV)},MATCH({F_CAT}{r},{fc(FC_NAME)},0)),0))'
        ws[f'{F_DACC}{r}'] = f'=IF({F_DEPT}{r}="","",IFERROR({code_of(f"INDEX({dp(DP_DEP)},MATCH({F_DEPT}{r},{dp(DP_NAME)},0))")},""))'
        ok = f'AND(ISNUMBER({F_COST}{r}),ISNUMBER({F_DATE}{r}),N({F_YEFF}{r})>0)'
        ws[f'{F_MDEP}{r}'] = f'=IF({ok},ROUND({F_COST}{r}*(1-N({F_SEFF}{r}))/({F_YEFF}{r}*12),2),"")'
        ws[f'{F_START}{r}'] = f'=IF(ISNUMBER({F_DATE}{r}),DATE(YEAR({F_DATE}{r}),MONTH({F_DATE}{r})+1,1),"")'
        ws[f'{F_END}{r}'] = f'=IF({ok},EDATE({F_START}{r},{F_YEFF}{r}*12-1),"")'
        ws[f'{F_LASTAMT}{r}'] = (f'=IF({ok},ROUND({F_COST}{r}*(1-N({F_SEFF}{r})),2)-{F_MDEP}{r}*({F_YEFF}{r}*12-1),"")')
        ws[f'{F_MONTHS}{r}'] = (f'=IF({F_MDEP}{r}="","",MAX(0,MIN({F_YEFF}{r}*12,(YEAR({rep_end})-YEAR({F_START}{r}))*12'
                                f'+MONTH({rep_end})-MONTH({F_START}{r})+1)))')
        ws[f'{F_ACCUM}{r}'] = (f'=IF({F_MDEP}{r}="","",IF({F_MONTHS}{r}>={F_YEFF}{r}*12,ROUND({F_COST}{r}*(1-N({F_SEFF}{r})),2),'
                               f'{F_MDEP}{r}*{F_MONTHS}{r}))')
        ws[f'{F_NBV}{r}'] = f'=IF({F_COST}{r}="","",{F_COST}{r}-N({F_ACCUM}{r}))'
        ws[f'{F_CHK}{r}'] = _chain(blank, [
            (f'{F_CAT}{r}=""', '✗ 没选类别'), (f'NOT(ISNUMBER({F_COST}{r}))', '✗ 原值没填或不是数字'),
            (f'NOT(ISNUMBER({F_DATE}{r}))', '✗ 没填入账日期'), (f'{F_DEPT}{r}=""', '✗ 没选使用部门（不知道折旧记到哪）'),
            (f'{F_DACC}{r}=""', '✗ 这个部门没设折旧科目'),
            (f'IFERROR(MATCH({F_DEPT}{r},{dp(DP_NAME)},0),99)>{N_DEPT_FA}', f'✗ 折旧只按【记账规则·部门】前 {N_DEPT_FA} 个部门算，这个部门排太后了'),
            (f'AND({F_YEARS}{r}<>"",NOT(ISNUMBER({F_YEARS}{r})))', '✗ 年限要填数字（比如 5，不要写「5年」）'),
            (f'AND({F_SALV}{r}<>"",NOT(ISNUMBER({F_SALV}{r})))', '✗ 残值率要填数字（比如 5%）'),
            (f'N({F_YEFF}{r})<=0', '✗ 没有折旧年限（这个类别也没设默认年限）'),
            (f'IF(ISNUMBER({F_YEFF}{r}),{F_YEFF}{r}*12<>INT({F_YEFF}{r}*12),FALSE)', '✗ 年限×12 要是整月数')])
        ws[f'{F_OK}{r}'] = f'=IF(AND(NOT({blank}),{F_CHK}{r}=""),1,0)'
    style_rows(ws, MOD_R0, FA_R1, 'ABCDEFGHIJKLMNOPQ', auto=(F_ACC, F_DACC, F_MDEP, F_START, F_MONTHS, F_ACCUM, F_NBV, F_CHK),
               fmts={F_COST: MONEY, F_DATE: DATE, F_YEARS: '0', F_SALV: '0%', F_MDEP: MONEY, F_START: 'yyyy/mm',
                     F_MONTHS: '0', F_ACCUM: MONEY, F_NBV: MONEY},
               aligns={F_NAME: AL, F_REMK: AL, F_CHK: AL, F_COST: AR_, F_MDEP: AR_, F_ACCUM: AR_, F_NBV: AR_})
    for r in range(MOD_R0, FA_R1 + 1):
        for c in (F_YEFF, F_SEFF, F_END, F_LASTAMT, F_OK):
            ws[f'{c}{r}'].font = F_NOTE
        ws[f'{F_END}{r}'].number_format = 'yyyy/mm'
        ws[f'{F_LASTAMT}{r}'].number_format = MONEY
    hide(ws, [F_YEFF, F_SEFF, F_END, F_LASTAMT, F_OK])
    _chk_fmt(ws, F_CHK, MOD_R0, FA_R1)
    add_date_dv(ws, f'{F_DATE}{MOD_R0}:{F_DATE}{FA_R1}')
    add_list_dv(ws, f'{F_CAT}{MOD_R0}:{F_CAT}{FA_R1}', '=资产类别列表')
    add_list_dv(ws, f'{F_DEPT}{MOD_R0}:{F_DEPT}{FA_R1}', '=部门列表')
    ws.auto_filter.ref = f'A{MOD_HDR}:{F_CHK}{FA_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws


# ───────────────────────── 手工凭证 ─────────────────────────
def build_man(wb, ctx):
    ws = wb.create_sheet(SH_MAN)
    widths(ws, {MN_GRP: 8, MN_DATE: 11, MN_MEMO: 30, MN_CODE: 26, MN_NAME: 30, MN_DR: 16, MN_CR: 16, MN_CP: 12,
                MN_PRJ: 10, MN_ATT: 6, MN_CHK: 30, MN_SEQ: 5, MN_CODE2: 9, MN_HAS: 5, MN_VDATE: 10, MN_BAD: 5,
                MN_OK: 5, MN_VID: 8, MN_KEY: 8, MN_GKEY: 9})
    title(ws, '手工凭证（调账、计提、结转、分摊…多借多贷）', MN_CHK, C_MAN,
          '💡 一条分录一行；同一张凭证的几行「凭证组号」写同一个号（数字、文字都行，每个月可以从 1 重新编），借贷要平。科目从下拉选或直接填编码。'
          '同一张凭证的几行挨着写，日期填在第一行（后面几行不填就跟上一行）。一张凭证只要有一行有错（✗），整张先不记账，改好了自动记上。'
          '凭证号跟别的凭证一起按月自动排。')
    header(ws, MOD_HDR, [(MN_GRP, '凭证组号'), (MN_DATE, '日期'), (MN_MEMO, '摘要'), (MN_CODE, '科目（编码 名称）'),
                         (MN_NAME, '科目全称（自动）'), (MN_DR, '借方金额'), (MN_CR, '贷方金额'), (MN_CP, '往来单位'),
                         (MN_PRJ, '项目'), (MN_ATT, '附单\n据张'), (MN_CHK, '校验'), (MN_SEQ, '组内序'),
                         (MN_CODE2, '科目编码'), (MN_HAS, '有金额'), (MN_VDATE, '凭证日期'), (MN_BAD, '行有错'),
                         (MN_OK, '能记账'), (MN_VID, '凭证ID'), (MN_KEY, '组键'), (MN_GKEY, '组号+月')], C_MAN, height=36)
    rg = lambda col: mod(SH_MAN, col)
    G = rg(MN_GRP)
    GK = rg(MN_GKEY)
    DRr, CRr = rg(MN_DR), rg(MN_CR)
    for r in range(MOD_R0, MAN_R1 + 1):
        ws[f'{MN_CODE2}{r}'] = f'=IF({MN_CODE}{r}="","",{code_of(f"{MN_CODE}{r}")})'
        ws[f'{MN_NAME}{r}'] = (f'=IF({MN_CODE2}{r}="","",IFERROR(INDEX({acc(AC_FULL)},MATCH({MN_CODE2}{r},{CODES},0)),'
                               f'"（科目表里没有）"))')
        ws[f'{MN_HAS}{r}'] = (f'=IF(AND({MN_GRP}{r}<>"",OR(N({MN_DR}{r})<>0,N({MN_CR}{r})<>0,'
                              f'AND({MN_DR}{r}<>"",NOT(ISNUMBER({MN_DR}{r}))),AND({MN_CR}{r}<>"",NOT(ISNUMBER({MN_CR}{r}))))),1,0)')
        # 凭证日期：这一行填了日期就用它；没填就跟紧挨着的上一行（同一个组号）走——一张凭证的几行要挨着写
        own = f'IF({MN_DATE}{r}<>"",IF(ISNUMBER({MN_DATE}{r}),{MN_DATE}{r},""),'
        if r == MOD_R0:
            ws[f'{MN_VDATE}{r}'] = f'=IF({MN_GRP}{r}="","",{own}""))'
        else:
            ws[f'{MN_VDATE}{r}'] = (f'=IF({MN_GRP}{r}="","",{own}IF({MN_GRP}{r - 1}&""={MN_GRP}{r}&"",{MN_VDATE}{r - 1},"")))')
        # 凭证分组键＝组号＋月份：组号每个月可以从 1 重新编
        ws[f'{MN_GKEY}{r}'] = (f'=IF(OR({MN_GRP}{r}="",NOT(ISNUMBER({MN_VDATE}{r}))),"",'
                               f'{MN_GRP}{r}&"|"&(YEAR({MN_VDATE}{r})*100+MONTH({MN_VDATE}{r})))')
        num_ok = (f'AND(OR({MN_DR}{r}="",ISNUMBER({MN_DR}{r})),OR({MN_CR}{r}="",ISNUMBER({MN_CR}{r})))')
        ws[f'{MN_BAD}{r}'] = (f'=IF({MN_HAS}{r}=0,0,IFERROR(IF(AND(ISNUMBER({MN_VDATE}{r}),{MN_VDATE}{r}>=起始月份,'
                              f'{month_no(f"{MN_VDATE}{r}")}<={N_MONTHS},{num_ok},'
                              f'ISNUMBER(MATCH({MN_CODE2}{r},{CODES},0)),OR(N({MN_DR}{r})=0,N({MN_CR}{r})=0)),0,1),1))')
        diff = f'ROUND(SUMIF({GK},{MN_GKEY}{r},{DRr})-SUMIF({GK},{MN_GKEY}{r},{CRr}),2)'
        ws[f'{MN_OK}{r}'] = (f'=IF(OR({MN_HAS}{r}=0,{MN_BAD}{r}=1,{MN_GKEY}{r}=""),0,'
                             f'IF(AND(COUNTIFS({GK},{MN_GKEY}{r},{rg(MN_BAD)},1)=0,{diff}=0),1,0))')
        ws[f'{MN_SEQ}{r}'] = (f'=IF({MN_OK}{r}=0,"",COUNTIFS(${MN_GKEY}${MOD_R0}:{MN_GKEY}{r},{MN_GKEY}{r},'
                              f'${MN_OK}${MOD_R0}:{MN_OK}{r},1))')
        # 凭证ID ＝ 来源×100000 ＋ 这张凭证第一条能记账的分录所在的位置
        ws[f'{MN_KEY}{r}'] = f'=IF({MN_OK}{r}=0,"",{MN_GKEY}{r})'
        ws[f'{MN_VID}{r}'] = f'=IF({MN_OK}{r}=0,"",{SRC["手工"] * 100000}+MATCH({MN_GKEY}{r},{rg(MN_KEY)},0))'
        blank = f'AND({MN_GRP}{r}="",{MN_DATE}{r}="",{MN_CODE}{r}="",{MN_DR}{r}="",{MN_CR}{r}="")'
        ws[f'{MN_CHK}{r}'] = _chain(blank, [
            (f'AND({MN_CODE}{r}="",{MN_DR}{r}="",{MN_CR}{r}="")', ''),        # 只写了组号/日期/摘要的抬头行，不算错
            (f'{MN_GRP}{r}=""', '✗ 没填凭证组号'),
            (f'AND({MN_DATE}{r}<>"",NOT(ISNUMBER({MN_DATE}{r})))', '✗ 日期不是真日期（要像 2026/9/30 这样填）'),
            (f'NOT(ISNUMBER({MN_VDATE}{r}))', '✗ 没有日期：这张凭证的第一行要填日期（同一张凭证的几行要挨着写）'),
            (f'{MN_VDATE}{r}<起始月份', '✗ 日期早于建账起始月份'),
            (f'{month_no(f"{MN_VDATE}{r}")}>{N_MONTHS}', '✗ 超出模板月份范围'),
            (f'OR(AND({MN_DR}{r}<>"",NOT(ISNUMBER({MN_DR}{r}))),AND({MN_CR}{r}<>"",NOT(ISNUMBER({MN_CR}{r}))))', '✗ 金额要填数字'),
            (f'{MN_CODE2}{r}=""', '✗ 没填科目'), (f'ISNA(MATCH({MN_CODE2}{r},{CODES},0))', '✗ 科目不在科目表里'),
            (f'AND(N({MN_DR}{r})<>0,N({MN_CR}{r})<>0)', '✗ 一行只填借方或贷方'),
            (f'{diff}<>0', f'="✗ 这张凭证借贷不平，差 "&FIXED({diff},2)'),
            (f'COUNTIFS({GK},{MN_GKEY}{r},{rg(MN_BAD)},1)>0', '✗ 这张凭证别的行有错，整张先不记账'),
            (f'INDEX({acc(AC_LEAF)},MATCH({MN_CODE2}{r},{CODES},0))<>"是"', '⚠ 不是末级科目（最好记到明细）')])
    style_rows(ws, MOD_R0, MAN_R1, 'ABCDEFGHIJK', auto=(MN_NAME, MN_CHK),
               fmts={MN_GRP: '0', MN_DATE: DATE, MN_DR: MONEY, MN_CR: MONEY, MN_CODE: '@'},
               aligns={MN_MEMO: AL, MN_CODE: AL, MN_NAME: AL, MN_CHK: AL, MN_DR: AR_, MN_CR: AR_})
    for r in range(MOD_R0, MAN_R1 + 1):
        ws[f'{MN_VDATE}{r}'].number_format = DATE
    _group_style(ws, MOD_R0, MAN_R1, [MN_SEQ, MN_CODE2, MN_HAS, MN_VDATE, MN_BAD, MN_OK, MN_VID, MN_KEY, MN_GKEY])
    _chk_fmt(ws, MN_CHK, MOD_R0, MAN_R1)
    add_date_dv(ws, f'{MN_DATE}{MOD_R0}:{MN_DATE}{MAN_R1}')
    add_list_dv(ws, f'{MN_CODE}{MOD_R0}:{MN_CODE}{MAN_R1}', '=科目选择', '从下拉选，或直接填编码', stop=False)
    add_list_dv(ws, f'{MN_CP}{MOD_R0}:{MN_CP}{MAN_R1}', '=往来列表', stop=False)
    add_list_dv(ws, f'{MN_PRJ}{MOD_R0}:{MN_PRJ}{MAN_R1}', '=项目列表', stop=False)
    ws.auto_filter.ref = f'A{MOD_HDR}:{MN_CHK}{MAN_R1}'
    ws.freeze_panes = f'C{MOD_R0}'
    return ws
