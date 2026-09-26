# -*- coding: utf-8 -*-
"""【现金流水总表】—— 沿用原表2【1现金流水总表】A:S 的列序和公式（第 5 行表头、第 6 行起数据，行号跟原表一样），
公式里的名字（账户名、类别表、记账汇率区…）原来指向另一个文件，现在指向本册【基础资料】【记账规则】【月度汇率】。
右边加了 4 列：指定对方科目（选填）、实际对方科目（自动）、凭证合并号（选填）、附单据张数；再往右是隐藏的记账辅助列。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *
from s_base import code_of


def _rule(col, cat):
    return f'INDEX({rng(SH_RULE, col, RC_R0, RC_R1)},MATCH({cat},类别表,0))&""'


def build_cash(wb, ctx):
    ws = wb.create_sheet(SH_CASH)
    widths(ws, {K_DATE: 11.5, K_ACC: 20, K_BANK: 9, K_CUR: 6, K_CAT: 15, K_ATTR: 8, K_MEMO: 38, K_IN: 15, K_OUT: 15,
                K_XR: 9, K_RATE: 9, K_NET: 15, K_UZS: 17, K_BAL: 16, K_BBAL: 18, K_PRJ: 10, K_CP: 12, K_NOTE: 20,
                K_CHK: 30, K_OVR: 12, K_CTR: 30, K_MERGE: 8, K_ATT: 6,
                K_ACODE: 9, K_RAW: 9, K_CCODE: 9, K_DR: 14, K_CR: 14, K_VALID: 5, K_CFI: 5, K_GSEQ: 5, K_GHEAD: 6,
                K_VID: 9, K_VMEMO: 30})
    title(ws, '现金流水总表 · 流水录入（五个账户全量合并，本位币苏姆）', K_LAST, C_CASH,
          '💡 一笔一行照原来记：选账户名称 → 银行、币种、汇率自动带出来；收入和支出分两列，只填一个；有实际成交价（购汇、结汇）就填「实际成交汇率」。'
          '选了「收/支类别」，右边自动出记账科目（按【记账规则】），每一笔自动生成一张记账凭证。'
          '要改某一笔的科目：在「指定对方科目」填编码（或下拉选）。几笔想并成一张凭证：「凭证合并号」写同一个号（要在同一个月）。'
          '新记录往下面空行填，别在中间插行。')
    # 第 3 行：汇总
    ws.row_dimensions[3].height = 26
    for i in range(1, CI(K_LAST) + 1):
        put(ws, f'{CL(i)}3', None, fill_=FILL_KPI, border=False)
    kp = [('A', '笔数：', 'B', f'=COUNT({cash(K_UZS)})', '#,##0'),
          ('E', '五个账户账面苏姆合计：', 'G', f'=SUM({rng(SH_BASE, A_NOWB, ACC_R0, ACC_R1)})', MONEY),
          ('I', '核对有问题的笔数：', 'K', f'=COUNTIF({cash(K_CHK)},"✗*")&" ✗ / "&COUNTIF({cash(K_CHK)},"⚠*")&" ⚠"', None),
          ('M', '生成凭证的笔数：', 'N', f'=SUM({cash(K_VALID)})', '#,##0')]
    for lc, lab, vc, f, fmt in kp:
        put(ws, f'{lc}3', lab, F_KPI_L, FILL_KPI, align=AR_, border=False)
        put(ws, f'{vc}3', f, Font(name=YH, sz=12, bold=True, color='FFC00000'), FILL_KPI, fmt, AL, border=False)
    ws.merge_cells('E3:F3')
    ws.merge_cells('I3:J3')
    # 第 4 行：分组色带
    ws.row_dimensions[4].height = 24
    bands = [('A', 'G', '登 记 信 息', 'FF5B9BD5'), ('H', 'H', '＋ 收 入', C_INCOME), ('I', 'I', '＋ 支 出', C_EXPENSE),
             ('J', 'M', '汇率 · 折苏姆（自动）', 'FF5B9BD5'), ('N', 'O', '余额（自动）', 'FF5B9BD5'),
             ('P', 'R', '项目 · 往来 · 备注', 'FF7030A0'), ('S', 'S', '核 对', 'FF833C0C'),
             (K_OVR, K_ATT, '记 账（新加）', 'FF548235')]
    for c0, c1, t, col in bands:
        if c0 != c1:
            ws.merge_cells(f'{c0}4:{c1}4')
        put(ws, f'{c0}4', t, Font(name=YH, sz=12, bold=True, color='FFFFFFFF'), fill(col), align=AC)
        for cc in range(CI(c0) + 1, CI(c1) + 1):
            put(ws, f'{CL(cc)}4', None, fill_=fill(col))
    hdr = [(K_DATE, '日期 ★'), (K_ACC, '账户名称 ★'), (K_BANK, '银行'), (K_CUR, '币种'), (K_CAT, '收/支类别 ★'),
           (K_ATTR, '类别属性'), (K_MEMO, '摘要'), (K_IN, '收入金额(原币)'), (K_OUT, '支出金额(原币)'),
           (K_XR, '实际成交汇率\n(选填)'), (K_RATE, '应用汇率'), (K_NET, '原币净额'), (K_UZS, '折苏姆'),
           (K_BAL, '账户余额(原币)'), (K_BBAL, '账面苏姆余额'), (K_PRJ, '项目/业务'), (K_CP, '往来单位'),
           (K_NOTE, '备注'), (K_CHK, '核对'), (K_OVR, '指定对方科目\n(选填)'), (K_CTR, '实际对方科目（自动）'),
           (K_MERGE, '凭证\n合并号'), (K_ATT, '附单\n据张')]
    hfill = {K_IN: 'FFE2EFDA', K_OUT: 'FFFFF2CC', K_OVR: 'FFE2EFDA', K_MERGE: 'FFE2EFDA', K_ATT: 'FFE2EFDA'}
    for col, t in hdr:
        put(ws, f'{col}{CASH_HDR}', t, F_HDR_D, fill(hfill.get(col, 'FFDDEBF7')), align=ACW)
    for col, t in [(K_ACODE, '本账户科目'), (K_RAW, '规则科目'), (K_CCODE, '对方科目'), (K_DR, '本账户借'),
                   (K_CR, '本账户贷'), (K_VALID, '有效'), (K_CFI, '现金流'), (K_GSEQ, '组内序'), (K_GHEAD, '组首'),
                   (K_VID, '凭证ID'), (K_VMEMO, '凭证摘要')]:
        put(ws, f'{col}{CASH_HDR}', t, F_NOTE, FILL_AUTO, align=ACW)
    ws.row_dimensions[CASH_HDR].height = 36

    rows = {x['row']: x for x in ctx['ledger']}
    ovr = ctx['overrides']
    acc_name = f'{rng(SH_ACC, AC_FULL, AC_R0, AC_R1)}'
    codes = acc(AC_CODE)
    cap = rng(SH_BASE, CP_CAP, CP_R0, CP_R1)
    for r in range(CASH_R0, CASH_R1 + 1):
        d = rows.get(r, {})
        k = r - CASH_R0 + 1
        # 录入列
        ws[f'{K_DATE}{r}'] = d.get('date')
        ws[f'{K_ACC}{r}'] = d.get('account')
        ws[f'{K_CAT}{r}'] = d.get('category')
        ws[f'{K_MEMO}{r}'] = d.get('memo')
        ws[f'{K_IN}{r}'] = d.get('inc')
        ws[f'{K_OUT}{r}'] = d.get('exp')
        ws[f'{K_XR}{r}'] = d.get('xrate')
        ws[f'{K_PRJ}{r}'] = d.get('project')
        ws[f'{K_CP}{r}'] = d.get('cp')
        ws[f'{K_NOTE}{r}'] = d.get('note')
        if r in ovr:
            ws[f'{K_OVR}{r}'] = ovr[r][0]
        # 原表的公式（名字改指本册）
        ws[f'{K_BANK}{r}'] = f'=IF($B{r}="","",IFERROR(INDEX(账户银行,MATCH($B{r},账户名,0)),""))'
        ws[f'{K_CUR}{r}'] = f'=IF($B{r}="","",IFERROR(INDEX(账户币种,MATCH($B{r},账户名,0)),""))'
        ws[f'{K_ATTR}{r}'] = f'=IF($E{r}="","",IFERROR(INDEX(类别属性,MATCH($E{r},类别表,0)),""))'
        ws[f'{K_RATE}{r}'] = (f'=IF(OR($A{r}="",$B{r}=""),"",IF($J{r}<>"",$J{r},IF($D{r}=本位币,1,IFERROR(LOOKUP(2,1/((汇率月份<=$A{r})'
                              f'*(INDEX(记账汇率区,0,MATCH($D{r},汇率币种头,0))<>"")),INDEX(记账汇率区,0,MATCH($D{r},汇率币种头,0))),""))))')
        ws[f'{K_NET}{r}'] = f'=IF(AND($H{r}="",$I{r}=""),"",ROUND(N($H{r})-N($I{r}),2))'
        ws[f'{K_UZS}{r}'] = f'=IF(OR($L{r}="",$K{r}=""),"",ROUND($L{r}*$K{r},2))'
        ws[f'{K_BAL}{r}'] = f'=IF($B{r}="","",ROUND(SUMIF(账户名,$B{r},账户期初)+SUMIFS($L${CASH_R0}:L{r},$B${CASH_R0}:B{r},$B{r}),2))'
        ws[f'{K_BBAL}{r}'] = (f'=IF($B{r}="","",IFERROR(ROUND(INDEX(账户期初本位,MATCH($B{r},账户名,0))'
                              f'+SUMIFS($M${CASH_R0}:M{r},$B${CASH_R0}:B{r},$B{r}),2),""))')
        # 记账辅助
        ws[f'{K_ACODE}{r}'] = f'=IF($B{r}="","",IFERROR(INDEX(账户科目,MATCH($B{r},账户名,0))&"",""))'
        raw = (f'IF(AND($A{r}<启用日,{_rule(RC_HIST, f"$E{r}")}<>""),{_rule(RC_HIST, f"$E{r}")},'
               f'IF(N($L{r})>=0,{_rule(RC_IN, f"$E{r}")},{_rule(RC_OUT, f"$E{r}")}))')
        ws[f'{K_RAW}{r}'] = f'=IF(OR($B{r}="",$L{r}=""),"",IF($T{r}<>"",$T{r}&"",IFERROR({raw},"")))'
        capc = f'IFERROR(INDEX({cap},MATCH($Q{r},单位表,0))&"","")'
        ws[f'{K_CCODE}{r}'] = (f'=IF($Y{r}="","",IF($Y{r}="@资本",IF({capc}="","3001",{capc}),{code_of(f"$Y{r}")}))')
        head = f'INDEX($A${CASH_R0}:$A${CASH_R1},MATCH($V{r},$V${CASH_R0}:$V${CASH_R1},0))'
        # 有效＝核对列里那些 ✗ 都没有（✗ 的行不出凭证）
        ok = (f'AND(ISNUMBER($A{r}),$A{r}>=起始月份,ISNUMBER($M{r}),$M{r}<>0,$X{r}<>"",$Z{r}<>"",'
              f'$E{r}<>"",ISNUMBER(MATCH($E{r},类别表,0)),'
              f'ISNUMBER(MATCH($X{r},{codes},0)),ISNUMBER(MATCH($Z{r},{codes},0)),NOT(AND($H{r}<>"",$I{r}<>"")),'
              f'OR($H{r}="",ISNUMBER($H{r})),OR($I{r}="",ISNUMBER($I{r})),'
              f'IF($V{r}="",TRUE,IFERROR(YEAR($A{r})*100+MONTH($A{r})=YEAR({head})*100+MONTH({head}),FALSE)))')
        ws[f'{K_VALID}{r}'] = f'=IF({ok},1,0)'
        ws[f'{K_DR}{r}'] = f'=IF($AC{r}=1,MAX($M{r},0),0)'
        ws[f'{K_CR}{r}'] = f'=IF($AC{r}=1,MAX(-$M{r},0),0)'
        # 现金流量项目：收钱按「现金流量(收)」、付钱按「现金流量(付)」。
        # 指定了对方科目的：资金账户/待对冲（账户间转钱）→ 0；长期资产 → 投资活动；
        # 原来类别不算现金流（内部转账）而改成损益科目的（比如手续费）→ 其他经营活动；其余仍按类别
        rule_cf = (f'IFERROR(IF($M{r}>0,N(INDEX({rng(SH_RULE, RC_CFIN, RC_R0, RC_R1)},MATCH($E{r},类别表,0))),'
                   f'N(INDEX({rng(SH_RULE, RC_CFOUT, RC_R0, RC_R1)},MATCH($E{r},类别表,0)))),0)')
        cf = (f'IF($T{r}="",{rule_cf},IF(OR(LEFT($Z{r},4)="1001",LEFT($Z{r},4)="1002",LEFT($Z{r},4)="1012",'
              f'$Z{r}="122104",$Z{r}="{EXCH_CODE}"),0,IF(OR(LEFT($Z{r},2)="16",LEFT($Z{r},2)="17"),IF($M{r}>0,10,12),'
              f'IF(AND(LEFT($Z{r},1)="5",{rule_cf}=0),IF($M{r}>0,2,6),{rule_cf}))))')
        ws[f'{K_CFI}{r}'] = f'=IF($AC{r}=1,{cf},"")'
        ws[f'{K_GSEQ}{r}'] = f'=IF(OR($AC{r}=0,$V{r}=""),"",COUNTIFS($V${CASH_R0}:$V{r},$V{r},$AC${CASH_R0}:$AC{r},1))'
        ws[f'{K_GHEAD}{r}'] = f'=IF($AE{r}=1,$V{r}&"","")'
        ws[f'{K_VID}{r}'] = (f'=IF($AC{r}=0,"",IF($V{r}="",{SRC["流水"] * 100000}+ROW()-{CASH_R0 - 1},'
                             f'{SRC["流水"] * 100000}+IFERROR(MATCH($V{r}&"",$AF${CASH_R0}:$AF${CASH_R1},0),ROW()-{CASH_R0 - 1})))')
        ws[f'{K_VMEMO}{r}'] = (f'=IF($AC{r}=0,"",$G{r}&IF($D{r}=本位币,"","（"&$D{r}&" "&FIXED(ABS($L{r}),2)'
                               f'&"×"&$K{r}&"）"))')
        # 实际对方科目（显示给人看）
        ws[f'{K_CTR}{r}'] = f'=IF($Z{r}="","",$Z{r}&" "&IFERROR(INDEX({acc_name},MATCH($Z{r},{codes},0)),"（科目表里没有）"))'
        # 核对
        ws[f'{K_CHK}{r}'] = (
            f'=IF(AND($A{r}="",$B{r}="",$E{r}="",$H{r}="",$I{r}=""),"",'
            f'IF($B{r}="","✗ 没选账户名称",IF(ISNA(MATCH($B{r},账户名,0)),"✗ 这个账户不在【基础资料】的账户档案里",'
            f'IF($A{r}="","✗ 没填日期",IF(NOT(ISNUMBER($A{r})),"✗ 日期不是真日期（要像 2026/9/25 这样填）",'
            f'IF($A{r}<起始月份,"✗ 日期早于建账起始月份",'
            f'IF(AND($H{r}<>"",$I{r}<>""),"✗ 收入、支出只能填一个",'
            f'IF(OR(AND($H{r}<>"",NOT(ISNUMBER($H{r}))),AND($I{r}<>"",NOT(ISNUMBER($I{r})))),"✗ 金额要填数字（不要带文字）",'
            f'IF(AND($H{r}="",$I{r}=""),IF(ISNUMBER(FIND("【待补】",$R{r})),"⚠ 源数据不全（源表缺金额），请补充","⚠ 没填金额"),'
            f'IF($E{r}="","✗ 没选收/支类别",IF(ISNA(MATCH($E{r},类别表,0)),"✗ 这个类别不在【记账规则】里",'
            f'IF($K{r}="","✗ 这个币种没有汇率（去【月度汇率】填）",'
            f'IF($X{r}="","✗ 这个账户在【基础资料】里没设对应科目",'
            f'IF(OR($Z{r}="",ISNA(MATCH($Z{r},{codes},0))),"✗ 对方科目「"&$Z{r}&"」不在科目表里",'
            f'IF(AND($V{r}<>"",YEAR($A{r})*100+MONTH($A{r})<>IFERROR(YEAR({head})*100+MONTH({head}),0)),'
            f'"✗ 同一个凭证合并号的几笔要在同一个月",'
            f'IF(AND($Y{r}="@资本",{capc}=""),"⚠ 投资款没对上股东（往来单位），先记在 3001 实收资本",'
            f'IF(ISNUMBER(FIND("【待补】",$R{r})),"⚠ 源数据不全（源表缺摘要或用途），请补充",'
            f'IF(AND($Y{r}<>"@资本",IFERROR(INDEX({acc(AC_LEAF)},MATCH($Z{r},{codes},0)),"是")<>"是"),"⚠ 对方科目不是末级科目，最好记到明细",'
            f'IF(AND(N($N{r})<0,N($N{r})-N($L{r})>=0),"⚠ 这一笔之后该账户余额变成负数了，核一下（后面继续是负的不再提示）","")))))))))))))))))))')
    # 样式
    inp = [K_DATE, K_ACC, K_CAT, K_MEMO, K_IN, K_OUT, K_XR, K_PRJ, K_CP, K_NOTE, K_OVR, K_MERGE, K_ATT]
    auto = [K_BANK, K_CUR, K_ATTR, K_RATE, K_NET, K_UZS, K_BAL, K_BBAL, K_CHK, K_CTR]
    style_rows(ws, CASH_R0, CASH_R1, inp + auto, auto=tuple(auto),
               fmts={K_DATE: DATE, K_IN: MONEY0, K_OUT: MONEY0, K_XR: RATE, K_RATE: RATE, K_NET: MONEY, K_UZS: MONEY,
                     K_BAL: MONEY, K_BBAL: MONEY, K_OVR: '@'},
               aligns={K_MEMO: AL, K_NOTE: AL, K_CHK: AL, K_CTR: AL, K_IN: AR_, K_OUT: AR_, K_NET: AR_, K_UZS: AR_,
                       K_BAL: AR_, K_BBAL: AR_, K_XR: AR_, K_RATE: AR_})
    for r in range(CASH_R0, CASH_R1 + 1):
        ws[f'{K_IN}{r}'].fill = fill('FFF3FAEF')
        ws[f'{K_OUT}{r}'].fill = fill('FFFFF8E1')
        for col in (K_ACODE, K_RAW, K_CCODE, K_DR, K_CR, K_VALID, K_CFI, K_GSEQ, K_GHEAD, K_VID, K_VMEMO):
            ws[f'{col}{r}'].font = F_NOTE
    hide(ws, [K_ACODE, K_RAW, K_CCODE, K_DR, K_CR, K_VALID, K_CFI, K_GSEQ, K_GHEAD, K_VID, K_VMEMO])
    rngc = f'{K_CHK}{CASH_R0}:{K_CHK}{CASH_R1}'
    ws.conditional_formatting.add(rngc, FormulaRule(formula=[f'LEFT(${K_CHK}{CASH_R0},1)="✗"'], fill=FILL_WARN,
                                                    font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rngc, FormulaRule(formula=[f'LEFT(${K_CHK}{CASH_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'{K_OVR}{CASH_R0}:{K_OVR}{CASH_R1}',
                                  FormulaRule(formula=[f'${K_OVR}{CASH_R0}<>""'], fill=fill('FFE2EFDA')))
    add_date_dv(ws, f'{K_DATE}{CASH_R0}:{K_DATE}{CASH_R1}')
    add_list_dv(ws, f'{K_ACC}{CASH_R0}:{K_ACC}{CASH_R1}', '=账户列表')
    add_list_dv(ws, f'{K_CAT}{CASH_R0}:{K_CAT}{CASH_R1}', '=类别列表', '收支类别决定记到哪个科目（见【记账规则】）')
    add_list_dv(ws, f'{K_PRJ}{CASH_R0}:{K_PRJ}{CASH_R1}', '=项目列表', stop=False)
    add_list_dv(ws, f'{K_CP}{CASH_R0}:{K_CP}{CASH_R1}', '=往来列表', '客户、供应商、股东、借支的员工写这里', stop=False)
    add_list_dv(ws, f'{K_OVR}{CASH_R0}:{K_OVR}{CASH_R1}', '=科目选择', '留空＝按类别自动；填了就用你填的科目', stop=False)
    ws.auto_filter.ref = f'A{CASH_HDR}:{K_LAST}{CASH_R1}'
    ws.freeze_panes = f'C{CASH_R0}'
    # 逐行修正的说明写进批注（不改你原来的「备注」）
    from openpyxl.comments import Comment
    for r, (code, why) in ovr.items():
        c = ws[f'{K_OVR}{r}']
        c.comment = Comment(f'生成时填的：{why}', '模板')
    return ws
