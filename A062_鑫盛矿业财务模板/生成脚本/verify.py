# -*- coding: utf-8 -*-
"""独立复核：不看模板里的公式，直接用 Python 从原表2 的流水 + 记账规则 + 修正清单把账重算一遍，
跟 LibreOffice 算出来的值（_lo_copy.xlsx）逐项比。

跑法：python3 verify.py [_lo_copy.xlsx]"""
import collections
import datetime as dt
import glob
import os
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from common import *
import fixes
import build

OK, BAD = [], []


def chk(name, cond, detail=''):
    (OK if cond else BAD).append((name, detail))
    print(('  ✓ ' if cond else '  ✗ ') + name + (f'  —— {detail}' if detail else ''))


def r2(x):
    return float(round(x + 0.0, 2))


def close(a, b, tol=0.011):
    return abs((a or 0) - (b or 0)) <= tol


def main(path):
    ctx = build.prepare()
    wv = openpyxl.load_workbook(path, data_only=True)
    golive = fixes.GOLIVE
    rules = {r[0]: r for r in ctx['rules']}
    chart = {a['code'] for a in ctx['chart']}
    months = ctx['months']
    rates = {m.strftime('%Y%m'): v for m, v in ctx['params']['rates']}
    cur_of = {a['name']: a['cur'] for a in ctx['accounts']}
    acode = fixes.ACCOUNT_CODE
    ovr = ctx['overrides']
    cap = {c['name']: c.get('cap') for c in ctx['counterparties']}

    def rate_at(cur, d):
        if cur == 'UZS':
            return 1.0
        best = None
        for m in months:
            if m <= d and rates.get(m.strftime('%Y%m'), {}).get(cur) not in (None, ''):
                best = float(rates[m.strftime('%Y%m')][cur])
        return best

    def code_of(x):
        x = str(x)
        return x.split(' ')[0] if ' ' in x else x

    # ── 1. 流水逐行：折苏姆、对方科目、是否记账 ──
    print('【1】现金流水逐行重算')
    ws = wv[SH_CASH]
    legs = []                          # (date, code, dr, cr, src)
    valid_rows = 0
    mism = []
    cash_rows = []
    for x in ctx['ledger']:
        r = x['row']
        d, acct, cat = x['date'], x['account'], x['category']
        inc, exp = x['inc'], x['exp']
        if inc in (None, '') and exp in (None, ''):
            continue
        net = r2((inc or 0) - (exp or 0))
        cur = cur_of.get(acct)
        rt = float(x['xrate']) if x['xrate'] not in (None, '') else rate_at(cur, d)
        uzs = r2(net * rt)
        got = ws[f'{K_UZS}{r}'].value
        if not close(got, uzs):
            mism.append((r, got, uzs))
        if not cat or not isinstance(d, dt.datetime):
            continue
        if r in ovr:
            cc = ovr[r][0]
        else:
            rule = rules[cat]
            if d < golive and rule[4]:
                cc = rule[4]
            else:
                cc = rule[2] if net >= 0 else rule[3]
        if cc == '@资本':
            cc = cap.get(x['cp']) or '3001'
        cc = code_of(cc)
        ac = acode[acct]
        if uzs == 0 or cc not in chart:
            continue
        valid_rows += 1
        cash_rows.append((r, d, acct, cur, net, uzs))
        if uzs > 0:
            legs += [(d, ac, uzs, 0, '流水'), (d, cc, 0, uzs, '流水')]
        else:
            legs += [(d, cc, -uzs, 0, '流水'), (d, ac, 0, -uzs, '流水')]
        got_c = ws[f'{K_CCODE}{r}'].value
        if str(got_c) != cc:
            mism.append((r, 'code', got_c, cc))
    chk('流水折苏姆、对方科目逐行一致', not mism, f'{len(mism)} 处不一致：{mism[:5]}')
    got_valid = sum(1 for r in range(CASH_R0, CASH_R1 + 1) if ws[f'{K_VALID}{r}'].value == 1)
    chk('生成凭证的流水笔数', got_valid == valid_rows, f'模板 {got_valid} / 复算 {valid_rows}')

    # ── 2. 月末调汇（外币账户） ──
    print('【2】外币账户月末调汇')
    rep = dt.datetime(2026, 8, 31)
    fx_cum = 0.0
    n_rep = (rep.year - months[0].year) * 12 + rep.month - months[0].month + 1
    fx_net = collections.defaultdict(float)
    for acct, cur in cur_of.items():
        if cur == 'UZS':
            continue
        # 逐月：累计调汇 = 月末原币余额×当月汇率 − 流水累计苏姆（汇率空的月份不调）
        cum = 0.0
        for n in range(1, n_rep + 1):
            m1 = dt.datetime(months[n - 1].year + months[n - 1].month // 12, months[n - 1].month % 12 + 1, 1)
            bal = sum(net for (r, d, a, c, net, u) in cash_rows if a == acct and d < m1)
            book = sum(u for (r, d, a, c, net, u) in cash_rows if a == acct and d < m1)
            rt = rates.get(months[n - 1].strftime('%Y%m'), {}).get(cur)
            if rt in (None, ''):
                continue
            new = r2(r2(bal * float(rt)) - book - cum)
            if new:
                legs.append((m1 - dt.timedelta(days=1), acode[acct], max(new, 0), max(-new, 0), '调汇'))
                fx_net[m1 - dt.timedelta(days=1)] += new
            cum = r2(cum + new)
        fx_cum += cum
    for d, v in fx_net.items():
        v = r2(v)
        if v:
            legs.append((d, '560304', max(-v, 0), max(v, 0), '调汇'))
    print(f'  复算：到 2026-08 累计调汇收益 {fx_cum:,.2f}（之前算的是 32,069,542.78）')
    chk('累计调汇收益 = 32,069,542.78', close(fx_cum, 32069542.78, 0.05), f'{fx_cum:,.2f}')

    # ── 3. 科目余额（报表月 2026-08 月末）逐个科目比 ──
    print('【3】科目余额表（2026-08 月末）逐科目比')
    tb = wv[SH_TB]
    hy, hm = wv[SH_HOME]['C4'].value, wv[SH_HOME]['F4'].value
    chk('首页报表期间是 2026 年 8 月', (hy, hm) == (2026, 8), f'{hy}-{hm}')
    bal = collections.defaultdict(float)
    for d, c, dr, cr, s in legs:
        if d <= rep:
            bal[c] += dr - cr
    got = {}
    for r in range(TB_R0, TB_R0 + (AC_R1 - AC_R0) + 1):
        c = tb[f'{T_CODE}{r}'].value
        if c:
            got[str(c)] = tb[f'{T_CNET}{r}'].value or 0
    bad = []
    for c in chart:
        exp_ = r2(sum(v for k, v in bal.items() if k.startswith(c)))
        if not close(got.get(c, 0), exp_, 0.05):
            bad.append((c, got.get(c), exp_))
    chk('每个科目（含上级汇总）月末余额 = 复算', not bad, f'{len(bad)} 个不一致：{bad[:8]}')
    lvl1 = sum(v for k, v in got.items() if len(k) == 4)
    chk('一级科目月末净额合计 = 0（借贷平衡）', close(lvl1, 0, 0.05), f'{lvl1:,.2f}')

    # ── 4. 利润表 2026 年 1-8 月 ──
    print('【4】利润表 2026 年 1-8 月')
    y0 = dt.datetime(2026, 1, 1)
    act = collections.defaultdict(float)
    for d, c, dr, cr, s in legs:
        if y0 <= d <= rep:
            act[c] += dr - cr
    A = lambda pre: sum(v for k, v in act.items() if k.startswith(pre))
    pl = wv[SH_PL]
    rev = -(A('5001') + A('5051'))
    cost = A('5401') + A('5402')
    net = -sum(v for k, v in act.items() if k.startswith('5'))
    from s_reports import pl_row
    chk('营业收入 = 6,219,055,000（跟你 8 月报表一样）', close(pl[f'C{pl_row(1)}'].value, rev) and close(rev, 6219055000),
        f'模板 {pl[f"C{pl_row(1)}"].value:,.2f} / 复算 {rev:,.2f}')
    chk('营业成本 = 2,779,561,189.02（跟 8 月报表一样）', close(pl[f'C{pl_row(2)}'].value, cost) and close(cost, 2779561189.02),
        f'模板 {pl[f"C{pl_row(2)}"].value:,.2f} / 复算 {cost:,.2f}')
    chk('税金及附加 = 323,005,250', close(pl[f'C{pl_row(3)}'].value, 323005250), f'{pl[f"C{pl_row(3)}"].value:,.2f}')
    chk('净利润 = 复算', close(pl[f'C{pl_row(32)}'].value, net, 0.05), f'模板 {pl[f"C{pl_row(32)}"].value:,.2f} / 复算 {net:,.2f}')
    print(f'    管理费用 {pl[f"C{pl_row(14)}"].value:,.2f}　财务费用 {pl[f"C{pl_row(18)}"].value:,.2f}　'
          f'营业外收入 {pl[f"C{pl_row(22)}"].value:,.2f}　营业外支出 {pl[f"C{pl_row(24)}"].value:,.2f}')

    # ── 5. 报表勾稽 ──
    print('【5】报表勾稽')
    bs = wv[SH_BAL]
    chk('资产负债表平衡（期末、年初）', close(bs['C37'].value, bs['G37'].value, 0.05) and close(bs['D37'].value, bs['H37'].value, 0.05),
        f'资产 {bs["C37"].value:,.2f} / 负债权益 {bs["G37"].value:,.2f}')
    cash_exp = sum(bal[c] for c in bal if c[:4] in ('1001', '1002', '1012'))
    chk('货币资金 = 复算', close(bs['C6'].value, cash_exp, 0.05), f'{bs["C6"].value:,.2f} / {cash_exp:,.2f}')
    cf = wv[SH_CF]
    last = max(r for r in range(5, 40) if cf[f'A{r}'].value == '五、期末现金余额')
    chk('现金流量表期末 = 货币资金', close(cf[f'C{last}'].value, bs['C6'].value, 0.05))
    chk('现金流量表本月期末 = 货币资金', close(cf[f'D{last}'].value, bs['C6'].value, 0.05))

    # ── 6. 凭证 ──
    print('【6】凭证')
    je = wv[SH_JE]
    R1 = ctx.get('je_r1') or je.max_row
    per_v = collections.defaultdict(lambda: [0.0, 0.0])
    n_lines = 0
    fut = 0
    for r in range(JE_R0, je.max_row + 1):
        if je[f'{J_NZ}{r}'].value == 1 and je[f'{J_DATE}{r}'].value > rep:
            fut += 1                       # 报表月以后的（月度汇率里已经填了 9、10 月汇率，9、10 月末自动调汇）
            continue
        if je[f'{J_NZ}{r}'].value == 1:
            n_lines += 1
            v = je[f'{J_VID}{r}'].value
            per_v[v][0] += je[f'{J_DR}{r}'].value or 0
            per_v[v][1] += je[f'{J_CR}{r}'].value or 0
    unb = [v for v, (a, b) in per_v.items() if not close(a, b)]
    chk('每张凭证借贷平衡', not unb, f'{len(per_v)} 张凭证，不平 {len(unb)}：{unb[:5]}')
    exp_lines = sum(1 for l in legs if (l[2] or l[3]) and l[0] <= rep)
    chk('分录条数（到 2026-08）= 复算', n_lines == exp_lines, f'模板 {n_lines} / 复算 {exp_lines}；报表月以后还有 {fut} 条（9、10 月调汇）')
    vl = wv[SH_VLIST]
    nv = sum(1 for r in range(VL_R0, VL_R0 + VL_N) if isinstance(vl[f'A{r}'].value, (int, float)))
    aug = sum(1 for (r, d, a, c, n_, u) in cash_rows if d.year == 2026 and d.month == 8) + \
        (1 if any(l[4] == '调汇' and l[0].year == 2026 and l[0].month == 8 for l in legs) else 0)
    chk('8 月凭证张数 = 8 月流水笔数 + 调汇', nv == aug, f'模板 {nv} / 复算 {aug}')
    nums = [vl[f'A{r}'].value for r in range(VL_R0, VL_R0 + nv)]
    chk('凭证号 1..N 连续', nums == list(range(1, nv + 1)))
    dates = [vl[f'B{r}'].value for r in range(VL_R0, VL_R0 + nv)]
    chk('凭证按日期排', all(dates[i] <= dates[i + 1] for i in range(len(dates) - 1)))
    pg = sum(vl[f'I{r}'].value or 0 for r in range(VL_R0, VL_R0 + nv))
    vp = wv[SH_VPRT]
    printed = sum(1 for p in range(VP_PAGES) if vp[f'I{p * VP_ROWS + 1}'].value not in (None, ''))
    chk('打印页数 = 凭证汇总总页数', printed == pg, f'{printed} / {pg}')
    # 抽第 1 页、最后一页
    if nv:
        r = 1
        memo = vp[f'B{r + 3}'].value
        chk('打印第 1 页有分录', bool(memo), str(memo))
        hdr = vp[f'E{r + 1}'].value
        chk('凭证号默认不印（留空）', '　　　　' in (hdr or ''), hdr)

    # ── 7. 数据校验 ──
    print('【7】数据校验表')
    ck = wv[SH_CHK]
    for r in range(5, 60):
        v = ck[f'C{r}'].value
        if v is None:
            continue
        mark = str(v)[:1]
        print(f'    {ck[f"B{r}"].value}：{v}')
    print(f'\n合计：✓ {len(OK)}　✗ {len(BAD)}')
    return len(BAD)


if __name__ == '__main__':
    p = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '_lo_copy.xlsx')
    sys.exit(1 if main(p) else 0)
