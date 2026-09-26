# -*- coding: utf-8 -*-
"""从两个原表里把数据读出来。

原表2（鑫盛矿业）的【1现金流水总表】公式都指向另一个没给我们的文件（苏姆本位资金日记账…全量合并版.xlsx）
的【参数设置】【月度汇率】。那个文件没给，但原表2 里存着它的缓存值（xl/externalLinks/externalLink2.xml），
账户档案、38 个收支类别、月度汇率都能原样恢复出来——这里就是从缓存里读的。
"""
import datetime as dt
import re
import zipfile

import openpyxl
from lxml import etree

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


def _s(v):
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def excel_date(v):
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    if isinstance(v, (int, float)) and 20000 < v < 80000:
        return dt.datetime(1899, 12, 30) + dt.timedelta(days=float(v))
    return v


class Src2:
    """原表2：鑫盛矿业 现金流水总表 + 科目表"""

    def __init__(self, path):
        self.path = path
        self.wv = openpyxl.load_workbook(path, data_only=True)

    # ── 现金流水：第 5 行表头，第 6 行起数据；只读要填的列（其余是公式，新表里重新写公式）──
    def ledger(self):
        ws = self.wv['1现金流水总表']
        rows = []
        for r in range(6, ws.max_row + 1):
            v = {c: ws[f'{c}{r}'].value for c in 'ABEGHIJPQR'}
            if all(_s(x) is None for x in v.values()):
                continue
            rows.append(dict(row=r, date=excel_date(v['A']), account=_s(v['B']), category=_s(v['E']),
                             memo=_s(v['G']), inc=v['H'], exp=v['I'], xrate=v['J'],
                             project=_s(v['P']), cp=_s(v['Q']), note=_s(v['R'])))
        return rows

    # ── 科目表：编码 名称 类别 余额方向 辅助核算类别 外币核算 是否现金科目 状态 是否平行科目 ──
    def accounts(self):
        ws = self.wv['2 科目表']
        res = []
        for r in range(2, ws.max_row + 1):
            v = [ws.cell(r, c).value for c in range(1, 10)]
            if v[0] in (None, ''):
                continue
            code = str(v[0]).strip()
            res.append(dict(code=code, name=_s(v[1]), cls=_s(v[2]), dir=_s(v[3]), aux=_s(v[4]),
                            fx=_s(v[5]), cashf=_s(v[6]), stat=_s(v[7]), para=_s(v[8])))
        return res

    # ── 外部链接缓存：参数设置 / 月度汇率 ──
    def _ext_cells(self):
        z = zipfile.ZipFile(self.path)
        name = None
        for n in z.namelist():
            if n.startswith('xl/externalLinks/externalLink') and n.endswith('.xml'):
                x = z.read(n)
                if '参数设置'.encode() in x:
                    name = n
        x = etree.fromstring(z.read(name))
        sheets = [s.get('val') for s in x.iter('{%s}sheetName' % NS['m'])]
        out = {}
        for sd in x.iter('{%s}sheetData' % NS['m']):
            sh = sheets[int(sd.get('sheetId'))]
            for c in sd.iter('{%s}cell' % NS['m']):
                v = c.find('m:v', NS)
                if v is None:
                    continue
                t = v.text
                if c.get('t') not in ('str', 's') and t not in (None, ''):
                    try:
                        t = float(t)
                    except ValueError:
                        pass
                out[(sh, c.get('r'))] = t
        return out

    def params(self):
        e = self._ext_cells()
        g = lambda sh, a: e.get((sh, a))
        accounts = []
        for r in range(11, 42):
            name = g('参数设置', f'C{r}')
            if name:
                accounts.append(dict(name=name, bank=g('参数设置', f'B{r}'), cur=g('参数设置', f'E{r}'),
                                     open=g('参数设置', f'G{r}') or 0, openb=g('参数设置', f'I{r}') or 0))
        cats = []
        for r in range(11, 49):
            name = g('参数设置', f'O{r}')
            if name:
                cats.append((name, g('参数设置', f'P{r}')))
        curs = [g('月度汇率', f'{c}5') for c in 'BCDEFGHI']
        rates = []
        for r in range(6, 26):
            m = g('月度汇率', f'A{r}')
            if m in (None, ''):
                continue
            rates.append((excel_date(m), {curs[i]: g('月度汇率', f'{c}{r}') for i, c in enumerate('BCDEFGHI')
                                          if g('月度汇率', f'{c}{r}') not in (None, '')}))
        return dict(start=excel_date(g('参数设置', 'B3')), cur=g('参数设置', 'B4'), curname=g('参数设置', 'B8'),
                    accounts=accounts, categories=cats, rates=rates)


class Src1:
    """原表1：做帐管理系统（只用它的公司名、会计科目明细对照）"""

    def __init__(self, path):
        self.wv = openpyxl.load_workbook(path, data_only=True)

    def company(self):
        return _s(self.wv['基础信息']['D3'].value)
