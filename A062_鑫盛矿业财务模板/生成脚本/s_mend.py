# -*- coding: utf-8 -*-
"""月末自动计算：
- _辅助：物料 × 月份 的全月一次加权平均（领用出库的单价从这里取）
- 【原矿产销存】：每月出矿量、转入成本、每吨成本、销售吨数、销售成本、结存
- 【月末结转】：折旧（按部门）、外币账户调汇、生产成本结转原矿、原矿结转销售成本、外币兑换差额（可选）

这里的数一律直接对各录入表求和，不对【记账分录】求和（记账分录又要引用这里，会循环引用）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *
from s_base import code_of

CODES = acc(AC_CODE)

# ─────────────────── _辅助：物料加权平均网格 ───────────────────
INV_G0 = CI('E')          # 第 1 个月的第一列
INV_W = 6                 # 每月 6 列：入库数量 入库金额 出库数量 单价 期末数量 期末金额
INV_FIELDS = ['入库数量', '入库金额', '出库数量', '加权单价', '期末数量', '期末金额']
INV_R0 = MAT_R0           # 跟物料档案同一行（第 5 行起）
INV_R1 = MAT_R1
INV_LAST = CL(INV_G0 + N_MONTHS * INV_W - 1)


def inv_col(n, field):
    return CL(INV_G0 + (n - 1) * INV_W + INV_FIELDS.index(field))


def inv_unit_cost(pos, n):
    """第 pos 个物料、第 n 个月的加权单价（pos、n 可以是公式）"""
    return f"INDEX({q(SH_AUX)}!${CL(INV_G0)}${INV_R0}:${INV_LAST}${INV_R1},{pos},({n}-1)*{INV_W}+4)"


def inv_field(pos, n, field):
    return f"INDEX({q(SH_AUX)}!${CL(INV_G0)}${INV_R0}:${INV_LAST}${INV_R1},{pos},({n}-1)*{INV_W}+{INV_FIELDS.index(field) + 1})"


def build_inv_grid(wb, ctx):
    ws = wb[SH_AUX]
    put(ws, 'A1', '物料 × 月份 全月一次加权平均（领用出库的单价、存货收发存从这里取）', F_NOTE, border=False)
    for c, t in (('A', '物料'), ('B', '库存管理'), ('C', '期初数量'), ('D', '期初金额')):
        put(ws, f'{c}4', t, F_NOTE, border=False)
    mat = lambda col: rng(SH_BASE, col, MAT_R0, MAT_R1)
    buy = lambda col: mod(SH_BUY, col)
    iss = lambda col: mod(SH_ISS, col)
    for n in range(1, N_MONTHS + 1):
        c0 = inv_col(n, '入库数量')
        ws[f'{c0}2'] = f'={mstart(n)}'
        ws[f'{c0}2'].number_format = 'yyyy/mm'
        ws[f'{c0}3'] = f'={mstart(n + 1)}'
        ws[f'{c0}3'].number_format = 'yyyy/mm'
        for f in INV_FIELDS:
            ws[f'{inv_col(n, f)}4'] = f'{n}月·{f}'
    for i in range(INV_R1 - INV_R0 + 1):
        r = INV_R0 + i
        k = i + 1
        ws[f'A{r}'] = f'=IF({at(mat(M_NAME), k)}="","",{at(mat(M_NAME), k)}&"")'
        ws[f'B{r}'] = f'=IF($A{r}="","",{at(mat(M_STOCK), k)}&"")'
        ws[f'C{r}'] = f'=N({at(mat(M_Q0), k)})'
        ws[f'D{r}'] = f'=N({at(mat(M_A0), k)})'
        for n in range(1, N_MONTHS + 1):
            qi, ai, qo, u, qe, ae = (inv_col(n, f) for f in INV_FIELDS)
            m0, m1 = f'{inv_col(n, "入库数量")}$2', f'{inv_col(n, "入库数量")}$3'
            if n == 1:
                qe_p, ae_p, u_p = f'$C{r}', f'$D{r}', f'IF($C{r}>0,$D{r}/$C{r},0)'
            else:
                qe_p, ae_p, u_p = (f'{inv_col(n - 1, "期末数量")}{r}', f'{inv_col(n - 1, "期末金额")}{r}',
                                   f'{inv_col(n - 1, "加权单价")}{r}')
            live = f'$B{r}="是"'
            # 进货只算能记账的行（有效=1，跟凭证一致）；领用只能用录入列（领用的单价从这里取，用它的「有效」会循环引用）
            ws[f'{qi}{r}'] = (f'=IF({live},SUMIFS({buy(B_QTY)},{buy(B_MAT)},$A{r},{buy(B_VALID)},1,{buy(B_DATE)},">="&{m0},'
                              f'{buy(B_DATE)},"<"&{m1}),0)')
            ws[f'{ai}{r}'] = (f'=IF({live},SUMIFS({buy(B_COST)},{buy(B_MAT)},$A{r},{buy(B_VALID)},1,{buy(B_DATE)},">="&{m0},'
                              f'{buy(B_DATE)},"<"&{m1}),0)')
            ws[f'{qo}{r}'] = (f'=IF({live},SUMIFS({iss(I_QTY)},{iss(I_MAT)},$A{r},{iss(I_USE)},"<>",{iss(I_DATE)},">="&{m0},'
                              f'{iss(I_DATE)},"<"&{m1}),0)')
            ws[f'{u}{r}'] = f'=IF({qe_p}+{qi}{r}>0,ROUND(({ae_p}+{ai}{r})/({qe_p}+{qi}{r}),6),{u_p})'
            ws[f'{qe}{r}'] = f'={qe_p}+{qi}{r}-{qo}{r}'
            ws[f'{ae}{r}'] = f'=ROUND({ae_p}+{ai}{r}-ROUND({qo}{r}*{u}{r},2),2)'
    return ws


# ─────────────────── 原矿产销存 ───────────────────
ORE_R0 = 6
ORE_R1 = ORE_R0 + N_MONTHS - 1
(O_MON, O_T0, O_A0, O_TIN, O_AIN, O_TAV, O_AAV, O_UC, O_TS, O_AS, O_TE, O_AE, O_DEV, O_GRADE, O_REV) = \
    'A B C D E F G H I J K L M N O'.split()


def ore_unit_cost(n):
    return f"INDEX({rng(SH_ORE, O_UC, ORE_R0, ORE_R1)},{n})"


def build_ore(wb, ctx):
    ws = wb.create_sheet(SH_ORE)
    widths(ws, {O_MON: 11, O_T0: 12, O_A0: 16, O_TIN: 12, O_AIN: 17, O_TAV: 12, O_AAV: 17, O_UC: 14, O_TS: 12,
                O_AS: 17, O_TE: 12, O_AE: 17, O_DEV: 11, O_GRADE: 9, O_REV: 17})
    title(ws, '原矿产销存（吨 · 苏姆）· 全自动', O_REV, C_PROD,
          '💡 出矿量取【产量登记】，转入成本＝当月结转的生产成本（【月末结转】），每吨成本＝（上月结存金额＋本月转入）÷（上月结存吨＋本月出矿吨），'
          '销售吨数取【销售结算】，销售成本＝销售吨×每吨成本（月末出一张结转销售成本的凭证）。还没出矿的月份都是 0。')
    header(ws, 5, [(O_MON, '月份'), (O_T0, '期初(吨)'), (O_A0, '期初金额'), (O_TIN, '本月出矿(吨)'),
                   (O_AIN, '本月转入成本'), (O_TAV, '可供(吨)'), (O_AAV, '可供金额'), (O_UC, '每吨成本'),
                   (O_TS, '销售(吨)'), (O_AS, '销售成本'), (O_TE, '期末(吨)'), (O_AE, '期末金额'),
                   (O_DEV, '掘进(米)'), (O_GRADE, '平均品位\n(%)'), (O_REV, '销售收入\n(不含税)')], C_PROD, height=36)
    pd = lambda col: mod(SH_PROD, col)
    sl = lambda col: mod(SH_SALE, col)
    for n in range(1, N_MONTHS + 1):
        r = ORE_R0 + n - 1
        m0, m1 = mstart(n), mstart(n + 1)
        ws[f'{O_MON}{r}'] = f'={m0}'
        dr = lambda c: f'{c},">="&{m0},{c},"<"&{m1}'
        ws[f'{O_T0}{r}'] = '=0' if n == 1 else f'={O_TE}{r - 1}'
        ws[f'{O_A0}{r}'] = '=0' if n == 1 else f'={O_AE}{r - 1}'
        ws[f'{O_TIN}{r}'] = f'=SUMIFS({pd(PD_TON)},{dr(pd(PD_DATE))})'
        ws[f'{O_AIN}{r}'] = f'={q(SH_MEND)}!{ME_SUM["cost"]}{ME_M0 + n - 1}'
        ws[f'{O_TAV}{r}'] = f'={O_T0}{r}+{O_TIN}{r}'
        ws[f'{O_AAV}{r}'] = f'={O_A0}{r}+{O_AIN}{r}'
        ws[f'{O_UC}{r}'] = f'=IF({O_TAV}{r}>0,ROUND({O_AAV}{r}/{O_TAV}{r},4),{0 if n == 1 else f"{O_UC}{r - 1}"})'
        ws[f'{O_TS}{r}'] = f'=SUMIFS({sl(S_TON)},{sl(S_VALID)},1,{dr(sl(S_DATE))})'     # 只算能记账的销售行
        ws[f'{O_AS}{r}'] = f'=IF({O_TS}{r}>={O_TAV}{r},{O_AAV}{r},ROUND({O_TS}{r}*{O_UC}{r},2))'
        ws[f'{O_TE}{r}'] = f'={O_TAV}{r}-{O_TS}{r}'
        ws[f'{O_AE}{r}'] = f'={O_AAV}{r}-{O_AS}{r}'
        ws[f'{O_DEV}{r}'] = f'=SUMIFS({pd(PD_DEV)},{dr(pd(PD_DATE))})'
        ws[f'{O_GRADE}{r}'] = (f'=IFERROR(SUMPRODUCT(({pd(PD_DATE)}>={m0})*({pd(PD_DATE)}<{m1})*ISNUMBER({pd(PD_GRADE)})'
                               f'*ISNUMBER({pd(PD_TON)}),{pd(PD_GRADE)},{pd(PD_TON)})/{O_TIN}{r},"")')
        ws[f'{O_REV}{r}'] = f'=SUMIFS({sl(S_NET)},{sl(S_VALID)},1,{dr(sl(S_DATE))})'
    style_rows(ws, ORE_R0, ORE_R1, 'ABCDEFGHIJKLMNO', bold=(O_UC, O_TE, O_AE),
               fmts={O_MON: MONTH, O_T0: QTY2, O_A0: MONEY, O_TIN: QTY2, O_AIN: MONEY, O_TAV: QTY2, O_AAV: MONEY,
                     O_UC: MONEY, O_TS: QTY2, O_AS: MONEY, O_TE: QTY2, O_AE: MONEY, O_DEV: '#,##0.0;-#,##0.0;"-"',
                     O_GRADE: '0.00', O_REV: MONEY},
               aligns={c: AR_ for c in 'BCDEFGHIJKLMNO'})
    r = ORE_R1 + 1
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for col in (O_TIN, O_AIN, O_TS, O_AS, O_DEV, O_REV):
        put(ws, f'{col}{r}', f'=SUM({col}{ORE_R0}:{col}{ORE_R1})', F_TXTB, FILL_TOT,
            QTY2 if col in (O_TIN, O_TS) else MONEY, AR_)
    for col in (O_T0, O_A0, O_TAV, O_AAV, O_UC, O_TE, O_AE, O_GRADE):
        put(ws, f'{col}{r}', None, F_TXTB, FILL_TOT)
    ws.conditional_formatting.add(f'{O_TE}{ORE_R0}:{O_TE}{ORE_R1}',
                                  FormulaRule(formula=[f'{O_TE}{ORE_R0}<0'], font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.freeze_panes = f'B{ORE_R0}'
    return ws


# ─────────────────── 月末结转 ───────────────────
ME_SUM = dict(mon='A', end='B', ton='C', dep='D', cost='E', ton_s='F', cogs='G', fx='H', xchg='I')
N_DEPT = 8                                   # 折旧：记账规则·部门的前 8 个
ME_DEP_HDR = ME_M1 + 4                       # 折旧块：部门名、折旧科目两行表头
ME_DEP0 = ME_DEP_HDR + 2
ME_DEP1 = ME_DEP0 + N_MONTHS - 1
ME_FXB0 = ME_DEP1 + 5                        # 调汇长表
ME_FXB1 = ME_FXB0 + N_ACC * N_MONTHS - 1
ME_X0 = ME_FXB1 + 5                          # 外币兑换差额（122105）
ME_X1 = ME_X0 + N_MONTHS - 1
ME_CCB0 = ME_X1 + 5                          # 生产成本结转长表（行数 = 成本科目数 × 月数）
DEP_COL = lambda j: CL(1 + j)                # 第 j 个部门在折旧块里的列（B..I）
FX_COLS = dict(n='A', acc='B', cur='C', code='D', bal='E', rate='F', should='G', book='H', prev='I', diff='J', cum='K')
CC_COLS = dict(n='A', code='B', amt='C', prev='D', carry='E', left='F')


def fx_row(i, n):            # 第 i 个账户（1 起）第 n 个月在调汇长表里的行
    return ME_FXB0 + (i - 1) * N_MONTHS + (n - 1)


def cc_row(j, n):            # 第 j 个成本科目第 n 个月
    return ME_CCB0 + (j - 1) * N_MONTHS + (n - 1)


N_CC_SLOTS = 20              # 生产成本末级科目最多 20 个（现在 11 个，科目表里新加的明细自动排进来）


def cost_leaves(ctx):
    """月末要结转到原矿的成本科目槽位：第 j 个 = 科目表里 4001/4101 下第 j 个末级科目（公式实时取，不是生成时定死）"""
    return list(range(1, N_CC_SLOTS + 1))


def cost_leaf_code(j):
    """第 j 个生产成本末级科目编码（科目表隐藏 P 列＝这类科目的排序键）"""
    p = rng(SH_ACC, 'P', AC_R0, AC_R1)
    return f'IFERROR(INDEX({acc(AC_CODE)},MATCH(SMALL({p},{j}),{p},0))&"","")'


def _src_sum_code(code, d0, d1):
    """各录入表里记到某科目的「借－贷」（日期在 [d0,d1) 内）。不含记账分录、月末结转本身（除了折旧，单独加）"""
    c = lambda col: cash(col)
    m = lambda sh, col: mod(sh, col)
    dr = lambda rg: f'{rg},">="&{d0},{rg},"<"&{d1}'
    parts = [
        # 现金流水：本账户科目那一边（借=收入）、对方科目那一边（借=本账户贷）
        f'SUMIFS({c(K_DR)},{c(K_ACODE)},{code},{dr(c(K_DATE))})-SUMIFS({c(K_CR)},{c(K_ACODE)},{code},{dr(c(K_DATE))})',
        f'SUMIFS({c(K_CR)},{c(K_CCODE)},{code},{dr(c(K_DATE))})-SUMIFS({c(K_DR)},{c(K_CCODE)},{code},{dr(c(K_DATE))})',
        # 采购：借 入账科目
        f'SUMIFS({m(SH_BUY, B_COST)},{m(SH_BUY, B_ACC)},{code},{m(SH_BUY, B_VALID)},1,{dr(m(SH_BUY, B_DATE))})',
        # 领用：借 计入科目、贷 存货科目
        f'SUMIFS({m(SH_ISS, I_AMT)},{m(SH_ISS, I_DR)},{code},{m(SH_ISS, I_VALID)},1,{dr(m(SH_ISS, I_DATE))})',
        f'-SUMIFS({m(SH_ISS, I_AMT)},{m(SH_ISS, I_CR)},{code},{m(SH_ISS, I_VALID)},1,{dr(m(SH_ISS, I_DATE))})',
        # 工资：借 计入、贷 应付
        f'SUMIFS({m(SH_PAY, W_AMT)},{m(SH_PAY, W_DR)},{code},{m(SH_PAY, W_VALID)},1,{dr(m(SH_PAY, W_DATE))})',
        f'-SUMIFS({m(SH_PAY, W_AMT)},{m(SH_PAY, W_CR)},{code},{m(SH_PAY, W_VALID)},1,{dr(m(SH_PAY, W_DATE))})',
        # 手工凭证（整张能记账的；日期按凭证第一行）
        f'SUMIFS({m(SH_MAN, MN_DR)},{m(SH_MAN, MN_CODE2)},{code},{m(SH_MAN, MN_OK)},1,{dr(m(SH_MAN, MN_VDATE))})'
        f'-SUMIFS({m(SH_MAN, MN_CR)},{m(SH_MAN, MN_CODE2)},{code},{m(SH_MAN, MN_OK)},1,{dr(m(SH_MAN, MN_VDATE))})',
    ]
    return '+'.join(parts).replace('+-', '-')


def build_mend(wb, ctx):
    ws = wb.create_sheet(SH_MEND)
    widths(ws, {'A': 11, 'B': 16, 'C': 14, 'D': 16, 'E': 17, 'F': 16, 'G': 17, 'H': 17, 'I': 17, 'J': 17, 'K': 17})
    title(ws, '月末结转（折旧 · 调汇 · 生产成本结转原矿 · 结转销售成本）· 全自动', 'K', C_JE,
          '💡 不用填，每个月自动算，凭证自动进【记账分录】（日期是月末最后一天）：'
          '① 折旧：【固定资产】卡片按部门汇总；② 调汇：人民币、美元账户按当月汇率重新折苏姆，差额记财务费用-汇兑损益；'
          '③ 生产成本结转：有出矿的月份，把生产成本（炸药、油料、配件、人工、折旧…）全部转成库存商品-原矿；'
          '④ 结转销售成本：卖了多少吨按每吨成本转主营业务成本；⑤ 外币兑换差额：【基础资料】参数选「是」才转汇兑损益。')
    # ── 每月汇总 ──
    put(ws, f'A{ME_M0 - 2}', '每月汇总（各结转凭证的金额）', F_SEC, border=False)
    header(ws, ME_M0 - 1, [('A', '月份'), ('B', '月末日期'), ('C', '出矿量(吨)'), ('D', '折旧合计'),
                           ('E', '生产成本结转\n原矿'), ('F', '销售(吨)'), ('G', '结转销售成本'), ('H', '调汇净收益\n(负＝损失)'),
                           ('I', '兑换差额转\n汇兑损益')], C_JE, height=36)
    leaves = cost_leaves(ctx)
    ctx['cost_leaves'] = leaves
    for n in range(1, N_MONTHS + 1):
        r = ME_M0 + n - 1
        ws[f'A{r}'] = f'={mstart(n)}'
        ws[f'B{r}'] = f'={mstart(n + 1)}-1'
        ws[f'C{r}'] = f"={q(SH_ORE)}!{O_TIN}{ORE_R0 + n - 1}"
        ws[f'D{r}'] = f'=ROUND(SUM({DEP_COL(1)}{ME_DEP0 + n - 1}:{DEP_COL(N_DEPT)}{ME_DEP0 + n - 1}),2)'
        ws[f'E{r}'] = '=ROUND(' + '+'.join(f'{CC_COLS["carry"]}{cc_row(j, n)}' for j in range(1, len(leaves) + 1)) + ',2)'
        ws[f'F{r}'] = f"={q(SH_ORE)}!{O_TS}{ORE_R0 + n - 1}"
        ws[f'G{r}'] = f"={q(SH_ORE)}!{O_AS}{ORE_R0 + n - 1}"
        ws[f'H{r}'] = '=ROUND(' + '+'.join(f'{FX_COLS["diff"]}{fx_row(i, n)}' for i in range(1, N_ACC + 1)) + ',2)'
        ws[f'I{r}'] = f'=D{ME_X0 + n - 1}'
    style_rows(ws, ME_M0, ME_M1, 'ABCDEFGHI', fmts={'A': MONTH, 'B': DATE, 'C': QTY2, 'D': MONEY, 'E': MONEY,
                                                   'F': QTY2, 'G': MONEY, 'H': MONEY, 'I': MONEY},
               aligns={c: AR_ for c in 'CDEFGHI'})

    # ── 折旧：按部门 ──
    put(ws, f'A{ME_DEP_HDR - 1}', '① 折旧（按【固定资产】卡片 · 使用部门）', F_SEC, border=False)
    dp = lambda col: rng(SH_RULE, col, DP_R0, DP_R1)
    fa = lambda col: mod(SH_FA, col)
    put(ws, f'A{ME_DEP_HDR}', '部门', F_HDR, fill(C_JE), align=AC)
    put(ws, f'A{ME_DEP_HDR + 1}', '折旧计入', F_HDR, fill(C_JE), align=AC)
    for j in range(1, N_DEPT + 1):
        c = DEP_COL(j)
        put(ws, f'{c}{ME_DEP_HDR}', f'=IF({at(dp(DP_NAME), j)}="","",{at(dp(DP_NAME), j)}&"")', F_HDR, fill(C_JE), align=AC)
        put(ws, f'{c}{ME_DEP_HDR + 1}', f'=IF({c}{ME_DEP_HDR}="","",{code_of(at(dp(DP_DEP), j))})', F_HDR_D, FILL_SUBH, align=AC)
    put(ws, f'J{ME_DEP_HDR}', '合计', F_HDR, fill(C_JE), align=AC)
    for n in range(1, N_MONTHS + 1):
        r = ME_DEP0 + n - 1
        m0 = f'$A{ME_M0 + n - 1}'
        ws[f'A{r}'] = f'={m0}'
        for j in range(1, N_DEPT + 1):
            c = DEP_COL(j)
            d = f'{c}${ME_DEP_HDR}'
            # 开始月 ≤ 本月 < 最后一个月：按月折旧额；最后一个月：用「末月折旧」补齐尾差（总额正好 = 原值×(1-残值率)）
            ws[f'{c}{r}'] = (f'=IF({d}="",0,ROUND(SUMIFS({fa(F_MDEP)},{fa(F_OK)},1,{fa(F_DEPT)},{d},{fa(F_START)},"<="&{m0},{fa(F_END)},">"&{m0})'
                             f'+SUMIFS({fa(F_LASTAMT)},{fa(F_OK)},1,{fa(F_DEPT)},{d},{fa(F_END)},{m0}),2))')
        ws[f'J{r}'] = f'=SUM(B{r}:{DEP_COL(N_DEPT)}{r})'
    style_rows(ws, ME_DEP0, ME_DEP1, 'ABCDEFGHIJ', fmts={'A': MONTH, **{CL(c): MONEY for c in range(2, 11)}},
               aligns={CL(c): AR_ for c in range(2, 11)})

    # ── 调汇长表 ──
    put(ws, f'A{ME_FXB0 - 2}', '② 外币账户调汇（每个资金账户 × 每月；苏姆账户差额恒为 0；只算能记账的流水）', F_SEC, border=False)
    header(ws, ME_FXB0 - 1, [('A', '月份'), ('B', '账户'), ('C', '币种'), ('D', '科目'), ('E', '月末原币余额'),
                             ('F', '当月汇率'), ('G', '应有苏姆'), ('H', '流水+手工累计苏姆'), ('I', '以前月份累计调汇'),
                             ('J', '本月调汇差额'), ('K', '累计调汇')], C_JE, height=36)
    ab = lambda col: rng(SH_BASE, col, ACC_R0, ACC_R1)
    for i in range(1, N_ACC + 1):
        for n in range(1, N_MONTHS + 1):
            r = fx_row(i, n)
            m1 = f'$B${ME_M0 + n - 1}+1'
            ws[f'A{r}'] = f'=$A${ME_M0 + n - 1}'
            ws[f'B{r}'] = f'=IF({at(ab(A_NAME), i)}="","",{at(ab(A_NAME), i)}&"")'
            ws[f'C{r}'] = f'=IF(B{r}="","",{at(ab(A_CUR), i)}&"")'
            ws[f'D{r}'] = f'=IF(B{r}="","",{at(ab(A_CODE), i)}&"")'
            ws[f'E{r}'] = (f'=IF(B{r}="",0,N({at(ab(A_OPEN), i)})+SUMIFS({cash(K_NET)},{cash(K_ACC)},B{r},{cash(K_VALID)},1,'
                           f'{cash(K_DATE)},"<"&{m1}))')
            # 汇率格子空着：Excel 里 INDEX 取空格子得 0，所以按数值判断（0 当没填）
            rt = f'INDEX(记账汇率区,{n},MATCH(C{r},汇率币种头,0))'
            ws[f'F{r}'] = f'=IF(OR(B{r}="",C{r}=本位币),1,IFERROR(IF(N({rt})=0,"",{rt}),""))'
            ws[f'G{r}'] = f'=IF(OR(B{r}="",F{r}=""),0,ROUND(E{r}*F{r},2))'
            # 账上已有的苏姆：只算记了账的流水（有效=1）+ 手工凭证 + 以前月份的调汇
            ws[f'H{r}'] = (f'=IF(B{r}="",0,N({at(ab(A_OPENB), i)})+SUMIFS({cash(K_UZS)},{cash(K_ACC)},B{r},{cash(K_VALID)},1,'
                           f'{cash(K_DATE)},"<"&{m1})'
                           f'+SUMIFS({mod(SH_MAN, MN_DR)},{mod(SH_MAN, MN_CODE2)},D{r},{mod(SH_MAN, MN_OK)},1,{mod(SH_MAN, MN_VDATE)},"<"&{m1})'
                           f'-SUMIFS({mod(SH_MAN, MN_CR)},{mod(SH_MAN, MN_CODE2)},D{r},{mod(SH_MAN, MN_OK)},1,{mod(SH_MAN, MN_VDATE)},"<"&{m1}))')
            ws[f'I{r}'] = '=0' if n == 1 else f'=K{r - 1}'
            ws[f'J{r}'] = f'=IF(OR(B{r}="",C{r}=本位币,F{r}=""),0,ROUND(G{r}-H{r}-I{r},2))'
            ws[f'K{r}'] = f'=I{r}+J{r}'
    style_rows(ws, ME_FXB0, ME_FXB1, 'ABCDEFGHIJK', fmts={'A': MONTH, 'E': MONEY, 'F': RATE, 'G': MONEY, 'H': MONEY,
                                                         'I': MONEY, 'J': MONEY, 'K': MONEY},
               aligns={c: AR_ for c in 'EFGHIJK'})

    # ── 外币兑换差额（122105） ──
    put(ws, f'A{ME_X0 - 2}', '⑤ 外币兑换待对冲（122105）余额转汇兑损益——参数「外币兑换差额月末转汇兑损益」为「是」，从「兑换差额从哪个月开始转」那个月起才转', F_SEC, border=False)
    header(ws, ME_X0 - 1, [('A', '月份'), ('B', '本月发生(借-贷)'), ('C', '上月末余额'), ('D', '本月转汇兑损益'),
                           ('E', '月末余额')], C_JE, height=30)
    for n in range(1, N_MONTHS + 1):
        r = ME_X0 + n - 1
        d0, d1 = f'$A${ME_M0 + n - 1}', f'($B${ME_M0 + n - 1}+1)'
        ws[f'A{r}'] = f'={d0}'
        ws[f'B{r}'] = f'=ROUND({_src_sum_code(chr(34) + "122105" + chr(34), d0, d1)},2)'
        ws[f'C{r}'] = '=0' if n == 1 else f'=E{r - 1}'
        # 从「兑换差额从哪个月开始转」那个月起才转；开始那个月把以前累计的一次转掉，以前的月份不动
        ws[f'D{r}'] = f'=IF(AND(兑换差额结转="是",ISNUMBER(兑换结转起始月),A{r}>=DATE(YEAR(兑换结转起始月),MONTH(兑换结转起始月),1)),C{r}+B{r},0)'
        ws[f'E{r}'] = f'=C{r}+B{r}-D{r}'
    style_rows(ws, ME_X0, ME_X1, 'ABCDE', fmts={'A': MONTH, 'B': MONEY, 'C': MONEY, 'D': MONEY, 'E': MONEY},
               aligns={c: AR_ for c in 'BCDE'})

    # ── 生产成本结转长表 ──
    put(ws, f'A{ME_CCB0 - 2}', f'③ 生产成本结转原矿（有出矿的月份，把生产成本各明细的余额全部转出；科目＝科目表里 4001/4101 下的末级科目，最多 {N_CC_SLOTS} 个，新加的自动排进来）', F_SEC, border=False)
    header(ws, ME_CCB0 - 1, [('A', '月份'), ('B', '成本科目'), ('C', '本月发生\n(借-贷)'), ('D', '上月未结转'),
                             ('E', '本月结转原矿'), ('F', '月末未结转')], C_JE, height=36)
    for j in leaves:
        for n in range(1, N_MONTHS + 1):
            r = cc_row(j, n)
            d0, d1 = f'$A${ME_M0 + n - 1}', f'($B${ME_M0 + n - 1}+1)'
            ws[f'A{r}'] = f'={d0}'
            ws[f'B{r}'] = f'={cost_leaf_code(j)}' if n == 1 else f'=B{cc_row(j, 1)}'
            dep = f'IF($B${ME_M0 + n - 1}<启用日,0,SUMIF(${DEP_COL(1)}${ME_DEP_HDR + 1}:${DEP_COL(N_DEPT)}${ME_DEP_HDR + 1},B{r},' \
                  f'${DEP_COL(1)}${ME_DEP0 + n - 1}:${DEP_COL(N_DEPT)}${ME_DEP0 + n - 1}))'
            ws[f'C{r}'] = f'=IF($B{r}="",0,ROUND({_src_sum_code(f"$B{r}", d0, d1)}+{dep},2))'
            ws[f'D{r}'] = '=0' if n == 1 else f'=F{r - 1}'
            ws[f'E{r}'] = f'=IF(N($C${ME_M0 + n - 1})>0,D{r}+C{r},0)'
            ws[f'F{r}'] = f'=D{r}+C{r}-E{r}'
    ctx['me_cc1'] = cc_row(len(leaves), N_MONTHS)
    style_rows(ws, ME_CCB0, ctx['me_cc1'], 'ABCDEF', fmts={'A': MONTH, 'C': MONEY, 'D': MONEY, 'E': MONEY, 'F': MONEY},
               aligns={c: AR_ for c in 'CDEF'})
    ws.freeze_panes = 'B6'
    return ws
