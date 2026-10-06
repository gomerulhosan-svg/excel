# -*- coding: utf-8 -*-
"""A050 1006 这一轮《03 财务账套与报表》的「账务骨架」改动（03A；底稿：参考/1006上传/03_财务账套与报表模板.xlsx）。

用法：apply(wb) 就地修改 openpyxl 工作簿（openpyxl.load_workbook(底稿)，不用 data_only），存盘后再调 post(path)。
      集成时先跑本模块，再跑 m1006_03b。对同一份底稿重复运行，结果完全一样（不读当前时间、不用随机数）。

这一轮改了什么（对应 spec 第 4 节 03A 和接口 K9～K12）：
 1. 新顺发泡网付款没进往来：【资金日记账】K 列「对应科目」，业务类型＝水果采购 / 包装物料采购 / 周转筐采购 的一律记
    应付账款（排在费用项目之前判断），付新顺、桂疆的 11 笔不再进管理费用，往来里冲掉应付。
 2. 对账补充行录了不进对账单：【往来业务明细】03补充段 K 列先认补充行 K 列（手工覆盖），再认 H 列（只要是数字），
    最后才用 数量×单价。【对账补充行】H 列改成蓝色录入格（空格里保留原公式）、L 列补录入色、第 2 行说明重写。
 3. 扩容（K10）：【资金日记账】扩到 4～4003（4,000 行），公式/格式/下拉/筛选跟着扩；全册 资金日记账!$603 → $4003；
    各汇总表的逐行辅助列（收入/支出/费用月度支出汇总、费用月度明细汇总、家用两张、国外费用汇总、借款汇总）扩到 4003 行，
    并改成按位置 INDEX 取（中间插行不再错位）；排序键 日期×1000 → ×10000；借款汇总「记账前借款」辅助段挪到 4004～4023；
    【_自动清单】往来单位池子接上日记账 604～4003 行（1926～5325），紧凑列表加护栏；「有收入的往来单位」池子 AF～AH
    扩到 4003 行、紧凑列 AJ 扩到 200 个（03B 果然鲜总表要用）；
    【记账凭证】在手工区（3010～3609）后面接「B区续」3610～10409 行，放日记账 604～4003 行的收付款凭证（k 接着编 601～4000）。
    收入/支出月度汇总 单位行 60 → 150（合计行挪到 157）；科目表费用项目 120 → 200 行（命名区跟着放宽），
    科目余额表/期初余额/记账凭证/各下拉 查 科目表!$A$4:$A$81 统一放宽到 $123。
 4. 往来（K9）：原 4～12304 段位置不动，只做 02筐子段 [2]周转筐出入库明细 O→P（单价）、W→X（备注）（配合《02》插装筐费列）、
    03补充段 K 列修正、03资金段改成按位置 INDEX 取；新段接在后面：
      12305～15704「03资金」续段（日记账 604～4003）；15705～16703「02筐子」续段（周转筐 2005～3003）；
      16704～17103「01果品」原料采购段（《01》【公司购买接口】4～403，应付卖方）；
      17104～17603「02装筐费」段（《02》【装筐费接口】4～503，应付装卸方）。全册 往来业务明细!$12304 → $17603。
 5. 物料销售成本（K11）：【对接源_02物料】W3:AB27 逐格读《02》【物料成本接口】A4:F27（24 个月）；
    【记账凭证】自动D区第一组按月生成「结转物料销售成本」：借 其他业务成本、贷 原材料—包装物料（科目取【收入成本映射】
    新加的「物料销售成本」一行，可改），跟结转开关无关，有成本就生成。
    顺带：【对接源_02物料】加「装筐费」一组（24 个月），【收入成本映射】加「装筐费」→ 借 主营业务成本、贷 应付账款，
    装筐费有了单价以后总账跟往来一起挂应付（单价默认空＝0，现在没有凭证）。
 6. 结转损益（K12）：新表【结转损益】：起始月 B2（默认 2026-08），12 个月每月一行：收入类合计、成本费用类合计、本月利润、
    物料销售成本、「是否结转」开关（下拉 是/否，默认全「否」）、提示（家用进费用、物料成本没取到、原料入库计价没填计价日期等）；
    下面按损益科目×月份列出净发生额（科目自动取科目表里类别＝损益的，最多 30 个）。
    【记账凭证】在「B区续」后面开「自动D区」：10410～10457 物料销售成本（24 个月×2 行），10458～10841 结转（12 个月×32 行：
    借各收入类科目、贷本年利润；借本年利润、贷各成本费用类科目），只有开关＝是 才出金额；合计行挪到 10842。
    结转表的取数范围只到 10457（含物料成本、不含结转凭证本身），不会循环引用。
    【利润表】改成按净额取（收入＝贷−借，费用＝借−贷）并排除结转凭证（O 列＝结转），利息收入 106.75 不再漏；
    【资产负债表】未分配利润始终并入「还没结转的损益余额」（已结转的月份损益余额是 0，不会重复），B3 改成公式自动显示
    「x 个月已结转」（取【结转损益】C18），原来的下拉删掉。
    【收入成本映射】里 3 个科目表没有的科目（主营业务收入—周转费、其他业务收入—筐子租赁费、其他业务收入—其他服务费）
    补进【科目表】69～71 行和【科目余额表】69～71 行（科目余额表 69～81 行公式统一成跟上面一样的按日期口径）。
 7. 【主页】加【结转损益】入口；【使用说明】补「每月怎么结转」「资金日记账只在最后一行下面接着录」等本轮要点。

 8. 复核后修正（同一轮）：
    ① 【结转损益】I 列提示每条单独一行（CHAR(10)），新增三条：【期初余额】原材料—包装物料还没登记《02》建账期初物料
       （金额取【对接源_02物料】W29:AB91 新块：逐格读《02》物料成本接口隐藏网格 60 种物料起始月的月初结存，不含筐；
       《02》网格表头变了 AB30 显示 ✘，提示改成「见《02》其他包装物料汇总」）；本月有水果销售收入、进销存对接「自购水果销售成本」
       为 0；本月有筐子销售收入、凭证没贷记周转材料—周转筐托盘。I:O 列宽放到 14/15、提示用 9 号字、行高按 6 条提示同时出现给够。
    ② 起始月锁定：B2 数据验证改成自定义公式，已有任一月选「是」就不能改（出错提示让满 12 个月另起下一年账套），G2、使用说明各补一句。
    ③ 【资产负债表】B3:C3 合并成公式格显示「x 个月已结转」，A3 文字改，原 C3 长说明删掉，删掉 B3 的下拉。
    ④ 设冻结前先清 sheet_view.selection（只有【结转损益】一处）。

post(path)：把《03》externalLink 里 WPS 写的绝对路径（/Users/Administrator/Desktop/…/02_物料与周转物台账模板.xlsx）
            改回相对路径 02_物料与周转物台账模板.xlsx。
"""
import os
import re
import sys
import copy
import shutil
import zipfile
import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.views import Selection
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment

from common0928 import (cell, title, note, section, header, widths, hide,
                        F_TITLE, F_NOTE, F_HDR, F_SEC, F_LBL, F_IN, F_AUTO, F_AUTOB, F_TOT, F_TOTN, F_HELP, F_CHK,
                        FL_TITLE, FL_HDR, FL_SEC, FL_AUTO, FL_IN, FL_LBL, FL_TOT, FL_NOTE, FL_NONE,
                        AC, AL, AR, BOX, NOB, MONEY, QTY, INT, DATE, MONTH, fill)

# ════════════════════════════════════════════════════════════════════════════
#  常量（跨本接口约定 K9～K12；03B 按这些行号写公式）
# ════════════════════════════════════════════════════════════════════════════
JR0, JR1_OLD, JR1 = 4, 603, 4003            # 资金日记账数据区
WL_OLD, WL_END = 12304, 17603               # 往来业务明细最后一行（旧 → 新）
SEG_ZJ = (12305, 15704)                     # 03资金续段 ← 日记账 604～4003
SEG_KZ = (15705, 16703)                     # 02筐子续段 ← [2]周转筐出入库明细 2005～3003
SEG_GG = (16704, 17103)                     # 01果品 原料采购 ← [1]公司购买接口 4～403
SEG_ZF = (17104, 17603)                     # 02装筐费 ← [2]装筐费接口 4～503
PZ_B0, PZ_B1 = 910, 2109                    # 记账凭证 B区（日记账 1～600）
PZ_M0, PZ_M1 = 3010, 3609                   # 手工区
PZ_X0 = 3610                                # B区续：日记账 601～4000（每笔 2 行）
PZ_X1 = PZ_X0 + 2 * (JR1 - JR1_OLD) - 1     # 10409
N_MC = 24                                   # 物料成本：24 个月 × 2 行
PZ_D0 = PZ_X1 + 1                           # 10410 自动D·物料成本
PZ_D1 = PZ_D0 + 2 * N_MC - 1                # 10457
N_ACC = 30                                  # 结转：每月 30 个损益科目位 ＋ 2 行本年利润
BLK = N_ACC + 2
PZ_J0 = PZ_D1 + 1                           # 10458 自动D·结转
PZ_J1 = PZ_J0 + 12 * BLK - 1                # 10841
PZ_END = PZ_J1                              # 凭证最后一行（合计之前）
PZ_TOT = PZ_END + 1                         # 10842 合计
KM_OLD, KM_NEW = 81, 123                    # 科目表 科目区 4～81 → 4～123
FY_OLD, FY_NEW = 123, 203                   # 科目表 费用项目 4～123 → 4～203（200 行）
POOL_OLD, POOL_NEW = 1925, 5325             # 03 _自动清单 往来单位池子
JZ = '结转损益'                              # K12
QC0, QC_N = 31, 60                          # 对接源_02物料 W31:Z90＝《02》物料成本接口 60 种物料起始月的月初结存
QC_TOT = QC0 + QC_N                         # 91：合计（不含筐）→【结转损益】期初物料提示

HELPER_SHEETS = ['收入月度汇总', '支出月度汇总', '费用月度支出汇总', '费用月度明细汇总',
                 '家用月度支出汇总', '家用月度明细汇总', '国外费用汇总', '借款汇总']


def _st(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)


def _ix(sheet, col, r0, r1, off):
    """按位置取：INDEX(表!$X$r0:$X$r1,ROW()-off)"""
    return f'INDEX({sheet}!${col}${r0}:${col}${r1},ROW()-{off})'


def _jx(col, off):
    return _ix('资金日记账', col, JR0, JR1, off)


def _tmpl(f, r):
    """把第 r 行的公式变成模板：本行行号 → {r}，上一行 → {r-1}（只认紧跟在列字母后面的相对行号）"""
    f = re.sub(r'(?<=[A-Z])%d(?!\d)' % r, '{r}', f)
    f = re.sub(r'(?<=[A-Z])%d(?!\d)' % (r - 1), '{r-1}', f)
    return f


def _fill(t, r):
    return t.replace('{r-1}', str(r - 1)).replace('{r-3}', str(r - 3)).replace('{r}', str(r))


def _units(s):
    """估算文字宽度（列宽单位）：中文/全角 2，其余 1.1"""
    return sum(2 if ord(ch) > 255 else 1.1 for ch in s)


def _lines(text, width, size=10):
    """按列宽 width（合并区各列宽之和）估算自动换行后的行数；\n 是硬换行"""
    k = size / 10
    return sum(max(1, -(-int(_units(p) * k) // int(width * 0.95))) for p in text.split('\n'))


def _each_formula(wb, sheets=None):
    for ws in wb.worksheets:
        if sheets is not None and ws.title not in sheets:
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    yield ws, c


# ════════════════════════════════════════════════════════════════════════════
#  一、全册区域改写（在加新行之前先把现有公式的写死范围放宽）
# ════════════════════════════════════════════════════════════════════════════
_RX = [
    # 资金日记账!$X$4:$X$603 / $1:$603 → 4003
    (re.compile(r'(资金日记账!\$[A-Z]{1,3}\$[14]:\$[A-Z]{1,3}\$)603(?!\d)'), r'\g<1>%d' % JR1),
    # 往来业务明细!$X$12304 / $X$4:$X$12304 / $1:$12304 → 17603
    (re.compile(r'(往来业务明细!\$[A-Z]{1,3}\$(?:\d+:\$[A-Z]{1,3}\$)?)%d(?!\d)' % WL_OLD), r'\g<1>%d' % WL_END),
    # 记账凭证!$X$4:$X$3609 → 凭证新最后一行
    (re.compile(r'(记账凭证!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)3609(?!\d)'), r'\g<1>%d' % PZ_END),
    # 科目表!$X$4:$X$81 → 123；科目表 费用项目/费用对应科目 $143 → 203
    (re.compile(r'(科目表!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)%d(?!\d)' % KM_OLD), r'\g<1>%d' % KM_NEW),
    (re.compile(r'(科目表!\$[MU]\$4:\$[MU]\$)143(?!\d)'), r'\g<1>%d' % FY_NEW),
]


def widen_all(wb):
    n = 0
    for ws, c in _each_formula(wb):
        v = c.value
        nv = v
        for rx, rep in _RX:
            nv = rx.sub(rep, nv)
        if ws.title == '往来业务明细':
            nv = re.sub(r'(\$[A-Z]{1,3}\$)%d(?!\d)' % WL_OLD, r'\g<1>%d' % WL_END, nv)
        if ws.title in HELPER_SHEETS:
            nv = re.sub(r'(\$[A-Z]{1,3}\$)603(?!\d)', r'\g<1>%d' % JR1, nv)
            if ws.title == '借款汇总':
                nv = re.sub(r'(\$[A-Z]{1,3}\$)623(?!\d)', r'\g<1>%d' % (JR1 + 20), nv)
            nv = nv.replace('*1000+', '*10000+').replace('/1000)*1000', '/10000)*10000')
        if nv != v:
            c.value = nv
            n += 1
    # 下拉里的科目表范围
    for ws in wb.worksheets:
        for dv in ws.data_validations.dataValidation:
            if dv.formula1 and '科目表!' in dv.formula1:
                dv.formula1 = re.sub(r'(科目表!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)%d(?!\d)' % KM_OLD,
                                     r'\g<1>%d' % KM_NEW, dv.formula1)
    # 定义名称：科目表 区域放宽（费用项目等 $Q$123 → $Q$203，科目名称/报表项目 $Q$81 → $Q$123）
    for name, dn in list(wb.defined_names.items()):
        t = dn.attr_text or ''
        if '科目表!' in t:
            t2 = t.replace('$Q$%d' % FY_OLD, '$Q$%d' % FY_NEW).replace('$Q$%d' % KM_OLD, '$Q$%d' % KM_NEW)
            if t2 != t:
                dn.attr_text = t2
    return n


# ════════════════════════════════════════════════════════════════════════════
#  二、资金日记账：K 列应付账款修正 ＋ 扩到 4003 行
# ════════════════════════════════════════════════════════════════════════════
def journal(wb):
    ws = wb['资金日记账']
    rx = re.compile(r'IF\(\$U(\d+)="是","应付账款",')
    n = 0
    for r in range(JR0, JR1_OLD + 1):
        v = ws[f'K{r}'].value
        nv = rx.sub(lambda m: (f'IF(OR($U{m.group(1)}="是",$F{m.group(1)}="水果采购",$F{m.group(1)}="包装物料采购",'
                               f'$F{m.group(1)}="周转筐采购"),"应付账款",'), v)
        if nv != v:
            ws[f'K{r}'] = nv
            n += 1
    assert n == JR1_OLD - JR0 + 1, n
    # 第 603 行当模板往下铺到 4003 行
    ncol = 33                                           # A～AG
    tm = {}
    for ci in range(1, ncol + 1):
        v = ws.cell(JR1_OLD, ci).value
        if isinstance(v, str) and v.startswith('='):
            tm[ci] = _tmpl(v, JR1_OLD)
    h = ws.row_dimensions[JR1_OLD].height
    src = [ws.cell(JR1_OLD, ci) for ci in range(1, ncol + 1)]
    for r in range(JR1_OLD + 1, JR1 + 1):
        for ci in range(1, ncol + 1):
            d = ws.cell(r, ci)
            _st(src[ci - 1], d)
            if ci in tm:
                d.value = _fill(tm[ci], r)
        if h:
            ws.row_dimensions[r].height = h
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(re.sub(r'(?<=[A-Z])%d(?!\d)' % JR1_OLD, str(JR1), str(dv.sqref)))
    ws.auto_filter.ref = f'A3:V{JR1}'
    a2 = ws['A2'].value or ''
    if '4,000' not in a2:
        ws['A2'] = ('★ 已扩到 4,000 行（第 4～4003 行）：只在最后一笔下面接着录，不要在中间插行（插行会让凭证、往来对不上）。'
                    + a2)


# ════════════════════════════════════════════════════════════════════════════
#  三、各汇总表的逐行辅助列扩到 4003 行（并改成按位置取）
# ════════════════════════════════════════════════════════════════════════════
_JREF = re.compile(r'资金日记账!\$([A-Z]{1,3})\{r\}')


def _helper_cols(ws, r0=JR0, r1=JR1_OLD):
    cols = []
    for ci in range(1, ws.max_column + 1):
        ok = True
        for r in range(r0, r1 + 1):
            v = ws.cell(r, ci).value
            if not (isinstance(v, str) and v.startswith('=')):
                ok = False
                break
        if ok:
            cols.append(ci)
    return cols


def _helper_tmpl(f, r):
    t = _tmpl(f, r)
    t = re.sub(r'=%d,1,0\)' % (r - 3), '={r-3},1,0)', t)
    t = re.sub(r'\*10000\+%d,' % r, '*10000+{r},', t)
    return _JREF.sub(lambda m: _jx(m.group(1), 3), t)


def helpers(wb):
    out = {}
    for name in HELPER_SHEETS:
        ws = wb[name]
        cols = _helper_cols(ws)
        out[name] = [L(c) for c in cols]
        if name == '借款汇总':
            continue
        for ci in cols:
            t = _helper_tmpl(ws.cell(JR1_OLD, ci).value, JR1_OLD)
            t4 = _fill(t, JR0)
            v4 = ws.cell(JR0, ci).value
            src = ws.cell(JR1_OLD, ci)
            for r in range(JR0, JR1 + 1):
                c = ws.cell(r, ci)
                if r == JR0 and v4 != t4:          # 第 4 行是累计列的起点（=Y4 这类），照原样只换取数写法
                    c.value = _JREF.sub(lambda m: _jx(m.group(1), 3), _tmpl(v4, JR0)).replace('{r}', str(JR0))
                else:
                    c.value = _fill(t, r)
                if r > JR1_OLD:
                    _st(src, c)
    _loans(wb)
    return out


def _loans(wb):
    """借款汇总：R～U 日记账段 4～4003、记账前借款段 4004～4023；V～AA 4～4023；AB 4～4003"""
    ws = wb['借款汇总']
    M0 = JR1 + 1                                        # 4004
    manual = {}
    for col in 'RSTU':
        manual[col] = [ws[f'{col}{r}'].value for r in range(604, 624)]
    sty = {col: [copy.copy(ws[f'{col}{r}']._style) for r in range(604, 624)] for col in 'RSTUVWXYZ'}
    sty['AA'] = [copy.copy(ws[f'AA{r}']._style) for r in range(604, 624)]
    # R～U：日记账段
    for col in 'RSTU':
        t = _helper_tmpl(ws[f'{col}603'].value, 603)
        src = ws[f'{col}603']
        for r in range(JR0, JR1 + 1):
            c = ws[f'{col}{r}']
            c.value = _fill(t, r)
            if r > JR1_OLD:
                _st(src, c)
        for i, v in enumerate(manual[col]):
            c = ws[f'{col}{M0 + i}']
            c.value = v
            c._style = sty[col][i]
    # V～AA：4～4023
    for col in ['V', 'W', 'X', 'Y', 'Z', 'AA']:
        t = _helper_tmpl(ws[f'{col}623'].value, 623)
        t4 = _fill(t, JR0)
        v4 = ws[f'{col}4'].value
        src = ws[f'{col}603']
        for r in range(JR0, M0 + 20):
            c = ws[f'{col}{r}']
            if r == JR0 and v4 != t4:
                continue
            c.value = _fill(t, r)
            if r > JR1_OLD:
                if r >= M0:
                    c._style = copy.copy(sty[col][r - M0])
                else:
                    _st(src, c)
    # AB：4～4003（排序键只排日记账段）
    t = _helper_tmpl(ws['AB603'].value, 603)
    src = ws['AB603']
    for r in range(JR0, JR1 + 1):
        c = ws[f'AB{r}']
        c.value = _fill(t, r)
        if r > JR1_OLD:
            _st(src, c)


# ════════════════════════════════════════════════════════════════════════════
#  四、收入 / 支出月度汇总：单位行 60 → 150
# ════════════════════════════════════════════════════════════════════════════
def monthly_rows(wb):
    for name in ('收入月度汇总', '支出月度汇总'):
        ws = wb[name]
        assert ws['A67'].value and '合' in ws['A67'].value, name
        tot = {c: (ws[f'{c}67'].value, copy.copy(ws[f'{c}67']._style)) for c in 'ABCDEFGHIJKLMN'}
        tm = {c: _tmpl(ws[f'{c}66'].value, 66) for c in 'ABCDEFGHIJKLMN'}
        h66, h67 = ws.row_dimensions[66].height, ws.row_dimensions[67].height
        for r in range(67, 157):
            for c in 'ABCDEFGHIJKLMN':
                d = ws[f'{c}{r}']
                d.value = _fill(tm[c], r)
                _st(ws[f'{c}66'], d)
            if h66:
                ws.row_dimensions[r].height = h66
        for c in 'ABCDEFGHIJKLMN':
            v, s = tot[c]
            d = ws[f'{c}157']
            d.value = v.replace('7:%s66)' % c, '7:%s156)' % c) if isinstance(v, str) else v
            d._style = s
        if h67:
            ws.row_dimensions[157].height = h67
        for c in 'BCDEFGHIJKLMN':
            v = ws[f'{c}5'].value
            if isinstance(v, str):
                ws[f'{c}5'] = v.replace(f'SUM({c}7:{c}66)', f'SUM({c}7:{c}156)')
        ws['D4'] = ws['D4'].value.replace('$N$67', '$N$157')


# ════════════════════════════════════════════════════════════════════════════
#  五、_自动清单：往来单位池子接上日记账 604～4003
# ════════════════════════════════════════════════════════════════════════════
def pool(wb):
    ws = wb['_自动清单']
    ix = lambda off: f'INDEX(资金日记账!$E$4:$E${JR1},ROW()-{off})'
    for r in range(306, 906):                           # 原来逐格 资金日记账!$E4 → 按位置取
        ws[f'A{r}'] = f'=IF({ix(305)}="","",{ix(305)})'
    src = [ws.cell(POOL_OLD, ci) for ci in (1, 2, 3)]
    for r in range(POOL_OLD + 1, POOL_NEW + 1):
        ws[f'A{r}'] = f'=IF({ix(POOL_OLD + 1 - 601)}="","",{ix(POOL_OLD + 1 - 601)})'
        ws[f'C{r}'] = f'=N(C{r - 1})+B{r}'
        for i, ci in enumerate((1, 2, 3)):
            _st(src[i], ws.cell(r, ci))
    for r in range(4, POOL_NEW + 1):
        ws[f'B{r}'] = f'=IF(A{r}="",0,IF(MATCH(A{r},$A$4:$A${POOL_NEW},0)=ROW()-3,1,0))'
    ws['A1'] = f'=N($C${POOL_NEW})'
    for r in range(4, 254):
        ws[f'E{r}'] = (f'=IF(ROW()-3>$A$1,"",IFERROR(INDEX($A$4:$A${POOL_NEW},'
                       f'MATCH(ROW()-3,$C$4:$C${POOL_NEW},0)),""))')
    # 「有收入的往来单位」池子 AF～AH（03B 果然鲜总表用紧凑列 AJ）：跟日记账一起扩到 4003 行，紧凑列 AJ 扩到 200 个
    src = [ws[f'{c}603'] for c in ('AF', 'AG', 'AH')]
    for r in range(JR0, JR1 + 1):
        e, h = _jx('E', 3), _jx('H', 3)
        ws[f'AF{r}'] = f'=IF(OR({e}="",N({h})=0),"",{e})'
        ws[f'AG{r}'] = f'=IF(AF{r}="",0,IF(MATCH(AF{r},$AF$4:$AF${JR1},0)=ROW()-3,1,0))'
        ws[f'AH{r}'] = f'=N(AH{r - 1})+AG{r}'
        if r > JR1_OLD:
            for i, c in enumerate(('AF', 'AG', 'AH')):
                _st(src[i], ws[f'{c}{r}'])
    ws['AF1'] = f'=MAX($AH$4:$AH${JR1})'
    for r in range(4, 204):
        ws[f'AJ{r}'] = (f'=IF(ROW()-3>$AH${JR1},"",IFERROR(INDEX($AF$4:$AF${JR1},'
                        f'MATCH(ROW()-3,$AH$4:$AH${JR1},0)),""))')
        _st(ws['AJ4'], ws[f'AJ{r}'])


# ════════════════════════════════════════════════════════════════════════════
#  六、记账凭证：B区续（日记账 601～4000）
# ════════════════════════════════════════════════════════════════════════════
def _vb(r, k, side):
    """B区一笔日记账的第 side（1/2）条分录"""
    J = lambda col: f'资金日记账!${col}$4:${col}${JR1}'
    e = 1 if side == 1 else 4
    return {
        'A': f'=IF($E{r}="","",ROW()-3)',
        'B': f'=IF($E{r}="","",INDEX({J("B")},{k}))',
        'C': f'=IF($E{r}="","","记-"&TEXT($B{r},"YYYYMM")&"-资金")',
        'D': f'=IF($E{r}="","",INDEX({J("D")},{k}))',
        'E': f'=IF(INDEX(资金日记账!$N$4:$S${JR1},{k},{e})=0,"",INDEX(资金日记账!$N$4:$S${JR1},{k},{e}))',
        'F': f'=IF($E{r}="","",IFERROR(INDEX(科目表!$A$4:$A${KM_NEW},MATCH($E{r},科目表!$B$4:$B${KM_NEW},0)),""))',
        'G': f'=IF($E{r}="","",IFERROR(INDEX(科目表!$E$4:$E${KM_NEW},MATCH($E{r},科目表!$B$4:$B${KM_NEW},0)),""))',
        'H': f'=IF($E{r}="","",N(INDEX(资金日记账!$N$4:$S${JR1},{k},{e + 1})))',
        'I': f'=IF($E{r}="","",N(INDEX(资金日记账!$N$4:$S${JR1},{k},{e + 2})))',
        'J': f'=IF($E{r}="","",INDEX({J("E")},{k}))',
        'K': f'=IF($E{r}="","",INDEX({J("F")},{k}))',
        'N': f'=IF($E{r}="","","系统")',
        'O': f'=IF($E{r}="","","自动B")',
    }


def vouchers_b(wb):
    ws = wb['记账凭证']
    # 先核对：用同一个写法生成第 910/911 行，应该跟现有公式一字不差
    for r, k, s in ((PZ_B0, 1, 1), (PZ_B0 + 1, 1, 2), (PZ_B1, 600, 2)):
        for col, f in _vb(r, k, s).items():
            assert ws[f'{col}{r}'].value == f, (r, col, ws[f'{col}{r}'].value, f)
    tot = {c: (ws[f'{c}3610'].value, copy.copy(ws[f'{c}3610']._style)) for c in 'ABCDEFGHIJKLMNOP'}
    sty = [{c: copy.copy(ws[f'{c}{PZ_B0 + i}']._style) for c in 'ABCDEFGHIJKLMNOP'} for i in (0, 1)]
    for k in range(JR1_OLD - 2, JR1 - 2):               # 601～4000
        for s in (1, 2):
            r = PZ_X0 + 2 * (k - (JR1_OLD - 2)) + (s - 1)
            for c in 'ABCDEFGHIJKLMNOP':
                ws[f'{c}{r}'].value = None
                ws[f'{c}{r}']._style = sty[s - 1][c]
            for col, f in _vb(r, k, s).items():
                ws[f'{col}{r}'] = f
    return tot


# ════════════════════════════════════════════════════════════════════════════
#  七、往来业务明细：O→P / W→X、03补充 K、03资金按位置取、四个新段
# ════════════════════════════════════════════════════════════════════════════
def _wl_tail(ws):
    """Q～X 列模板（从第 17603 行之前的老末行 12304 取；范围已经放宽到 17603）"""
    return {c: _tmpl(ws[f'{c}{WL_OLD}'].value, WL_OLD) for c in 'QRSTUVWX'}


def _seg_zj(r, off):
    I = lambda col: _jx(col, off)
    return {
        'C': f'=IFERROR(1*{I("B")},0)',
        'D': f'=IF(OR(E{r}="应收",E{r}="应付"),IF({I("E")}="","",{I("E")}),"")',
        'E': f'=IF({I("W")}="","",{I("W")})',
        'F': f'=IF({I("F")}="","",{I("F")})',
        'G': f'=IF({I("D")}="","",{I("D")})',
        'H': '', 'I': '', 'J': '',
        'K': f'=IF(E{r}="应收",IFERROR(1*{I("I")},0),IF(E{r}="应付",IFERROR(1*{I("H")},0),0))',
        'L': f'=IF(E{r}="应收",-IFERROR(1*{I("H")},0),IF(E{r}="应付",-IFERROR(1*{I("I")},0),0))',
        'M': f'=K{r}+L{r}',
        'N': f'=IF({I("L")}="","",{I("L")})',
        'O': f'="资金日记账第{r - off + 3}行"',
        'P': f'=IF({I("C")}="","",{I("C")})',
    }


def _seg_kz(r):
    off = SEG_KZ[0] - 2002                              # 15705 → 周转筐第 2005 行（区域 4～3003 的第 2002 个）
    X = lambda col: _ix('[2]周转筐出入库明细', col, 4, 3003, off)
    f = X('F')
    return {
        'C': f'=IFERROR(1*{X("B")},0)',
        'D': f'=IF(AND(E{r}<>"",K{r}<>0),IF({X("C")}="","",{X("C")}),"")',
        'E': (f'=IF(AND(OR({f}="采购入库",{f}="采购退货",{f}="销售出库",{f}="销售退回"),J{r}<>0,I{r}<>0),'
              f'IF(OR({f}="采购入库",{f}="采购退货"),"应付","应收"),"")'),
        'F': f'=IF(E{r}="","",IF(E{r}="应付","筐子托盘采购","筐子销售"))',
        'G': f'=IF({X("D")}="","",{X("D")})',
        'H': '个',
        'I': (f'=IF({f}="采购入库",IFERROR(1*{X("H")},0),IF({f}="采购退货",-IFERROR(1*{X("H")},0),'
              f'IF({f}="销售退回",-IFERROR(1*{X("I")},0),IFERROR(1*{X("K")},0))))'),
        'J': f'=IFERROR(1*{X("P")},0)',
        'K': f'=IF(E{r}="",0,ROUND(I{r}*J{r},2))',
        'L': 0,
        'M': f'=K{r}',
        'N': f'=IF({X("X")}="","",{X("X")})',
        'O': f'=IF({X("E")}="","",{X("E")})',
        'P': '',
    }


def _seg_gg(r):
    off = SEG_GG[0] - 1                                 # 16704 → 公司购买接口第 4 行
    G = lambda col: _ix('[1]公司购买接口', col, 4, 403, off)
    return {
        'C': f'=IFERROR(1*{G("B")},0)',
        'D': f'=IF(K{r}=0,"",IF({G("C")}="","",{G("C")}))',
        'E': '应付',
        'F': '果品采购（原料公司购买）',
        'G': f'=IF({G("D")}="","",{G("D")})',
        'H': f'=IF({G("G")}="","",{G("G")})',
        'I': f'=IFERROR(1*{G("F")},0)',
        'J': f'=IFERROR(1*{G("H")},0)',
        'K': f'=ROUND(IFERROR(1*{G("I")},0),2)',
        'L': 0,
        'M': f'=K{r}',
        'N': (f'=IF(K{r}=0,"","原料出库公司购买 "&{G("E")}&IF({G("K")}&""="","","；去向 "&{G("K")})'
              f'&IF({G("J")}&""="","","；"&{G("J")})&IF({G("N")}&""="","","；"&{G("N")}))'),
        'O': f'=IF({G("L")}="","",{G("L")})',
        'P': '',
    }


def _seg_zf(r):
    off = SEG_ZF[0] - 1                                 # 17104 → 装筐费接口第 4 行
    Z = lambda col: _ix('[2]装筐费接口', col, 4, 503, off)
    return {
        'C': f'=IFERROR(1*{Z("B")},0)',
        'D': f'=IF(K{r}=0,"",IF({Z("C")}="","",{Z("C")}))',
        'E': '应付',
        'F': '装筐费',
        'G': f'=IF({Z("D")}="","",{Z("D")})',
        'H': '个',
        'I': f'=IFERROR(1*{Z("E")},0)',
        'J': f'=IFERROR(1*{Z("F")},0)',
        'K': f'=ROUND(IFERROR(1*{Z("G")},0),2)',
        'L': 0,
        'M': f'=K{r}',
        'N': f'=IF({Z("I")}="","",{Z("I")})',
        'O': f'=IF(N({Z("H")})=0,"","周转筐第"&{Z("H")}&"行")',
        'P': '',
    }


def wanglai(wb):
    ws = wb['往来业务明细']
    # ① 02筐子段（2004～4004）：《02》周转筐出入库明细在 N 后插了 O「装筐费」，单价 O→P、备注 W→X
    n = 0
    for r in range(2004, 4005):
        for col in 'JN':
            v = ws[f'{col}{r}'].value
            nv = re.sub(r'(\[2\]周转筐出入库明细!)O(\d+)', r'\1P\2', v)
            nv = re.sub(r'(\[2\]周转筐出入库明细!)W(\d+)', r'\1X\2', nv)
            if nv != v:
                ws[f'{col}{r}'] = nv
                n += 1
    assert n == 2001 * 2, n
    # ② 03补充段 K：先认补充行 K（手工覆盖），再认 H（只要是数字），最后 数量×单价
    for r in range(12005, WL_OLD + 1):
        s = r - 12001
        ws[f'K{r}'] = (f'=IF(对账补充行!K{s}<>"",IFERROR(1*对账补充行!K{s},0),IF(ISNUMBER(对账补充行!H{s}),'
                       f'ROUND(对账补充行!H{s},2),ROUND(I{r}*J{r},2)))')
    # ③ 03资金段（11105～11704）改成按位置取
    for r in range(11105, 11705):
        assert ws[f'A{r}'].value == '03资金' and ws[f'B{r}'].value == r - 11101
        for col, f in _seg_zj(r, 11104).items():
            if col in 'HIJ':
                continue
            ws[f'{col}{r}'] = f
    # ④ 新段
    tail = _wl_tail(ws)
    segs = [
        (SEG_ZJ, '03资金', lambda r: r - (SEG_ZJ[0] - 604), lambda r: _seg_zj(r, SEG_ZJ[0] - 601), 11105),
        (SEG_KZ, '02筐子', lambda r: r - (SEG_KZ[0] - 2005), _seg_kz, 2005),
        (SEG_GG, '01果品', lambda r: r - (SEG_GG[0] - 4), _seg_gg, 8105),
        (SEG_ZF, '02装筐费', lambda r: r - (SEG_ZF[0] - 4), _seg_zf, 2005),
    ]
    for (r0, r1), name, srow, body, sty_row in segs:
        sty = {c: copy.copy(ws[f'{c}{sty_row}']._style) for c in 'ABCDEFGHIJKLMNOPQRSTUVWX'}
        h = ws.row_dimensions[sty_row].height
        for r in range(r0, r1 + 1):
            vals = {'A': name, 'B': srow(r)}
            vals.update(body(r))
            for c, t in tail.items():
                vals[c] = _fill(t, r)
            for c in 'ABCDEFGHIJKLMNOPQRSTUVWX':
                d = ws[f'{c}{r}']
                d.value = vals.get(c)
                d._style = sty[c]
            if h:
                ws.row_dimensions[r].height = h
    ws.auto_filter.ref = f'A3:P{WL_END}'


# ════════════════════════════════════════════════════════════════════════════
#  八、对账补充行界面
# ════════════════════════════════════════════════════════════════════════════
def buchong(wb):
    ws = wb['对账补充行']
    ws['H3'] = '金额（元）\n可直接填，负数＝冲减'
    ws['A2'] = ('金额直接填 H 列（也可以只填数量＋单价自动算，K 列可强制覆盖）；正数＝增加对方欠款（应收）或我方欠款（应付），'
                '负数＝冲减（抹零、优惠）；L 列选应收/应付（空白按应收）；I 列填已收/已付的正数。已记资金日记账的收付款不要重复补录。'
                '本表只进往来对账，不进记账凭证：抹零/优惠要进总账，请在记账凭证手工区补一笔。')
    ws.row_dimensions[2].height = 42
    ws['A2'].alignment = AL
    for r in range(4, 304):
        c = ws[f'H{r}']
        c.fill = fill('DDEBF7')
        c.font = Font(name='宋体', size=10, color='1F4E79')
        c = ws[f'L{r}']
        c.fill = fill('DDEBF7')
        c.font = Font(name='宋体', size=10, color='1F4E79')
        c.border = BOX
        c.alignment = AC
    ws['L3']._style = copy.copy(ws['K3']._style)
    ws.row_dimensions[3].height = max(ws.row_dimensions[3].height or 0, 30)
    ws.column_dimensions['H'].width = max(ws.column_dimensions['H'].width or 0, 16)
    ws.column_dimensions['L'].width = max(ws.column_dimensions['L'].width or 0, 14)


# ════════════════════════════════════════════════════════════════════════════
#  九、科目表 / 科目余额表 / 收入成本映射
# ════════════════════════════════════════════════════════════════════════════
NEW_ACC = [('5005', '主营业务收入—周转费'), ('5054', '其他业务收入—筐子租赁费'), ('5055', '其他业务收入—其他服务费')]


def kemu(wb):
    ws = wb['科目表']
    for i, (code, nm) in enumerate(NEW_ACC):
        r = 69 + i
        assert not ws[f'B{r}'].value
        for c, v in zip('ABCDEFG', (code, nm, '损益', '贷', '营业收入', '收入', '1006 补：【收入成本映射】用到的科目')):
            ws[f'{c}{r}'] = v
            _st(ws[f'{c}68'], ws[f'{c}{r}'])
    for r in range(82, KM_NEW + 1):
        for c in 'ABCDEFG':
            _st(ws[f'{c}81'], ws[f'{c}{r}'])
    ws.auto_filter.ref = f'A3:G{KM_NEW}'
    for r in range(116, FY_NEW + 1):                   # 费用项目 M、费用对应科目 U 录入格扩到 203 行
        for c in 'MU':
            _st(ws[f'{c}115'], ws[f'{c}{r}'])
    # 科目余额表：69～71 补科目名；69～81 公式统一成跟第 68 行一样的按日期口径
    kb = wb['科目余额表']
    for i, (code, nm) in enumerate(NEW_ACC):
        assert not kb[f'B{69 + i}'].value
        kb[f'B{69 + i}'] = nm
    tm = {c: _tmpl(kb[f'{c}68'].value, 68) for c in 'ACDEFGHIJ'}
    for r in range(69, KM_OLD + 1):
        for c, t in tm.items():
            kb[f'{c}{r}'] = _fill(t, r)
    # 收入成本映射：补「物料销售成本」「装筐费」两行
    mp = wb['收入成本映射']
    rows = [('25', '物料销售成本', '结转型', '0%', '其他业务成本', '原材料—包装物料',
             '《02》物料成本接口按月末加权算的物料销售成本；记账凭证自动D区每月自动结转'),
            ('26', '装筐费', '成本型', '0%', '主营业务成本', '应付账款',
             '《02》周转筐装筐费（单价在《02》基础资料），挂应付装卸方')]
    for i, vals in enumerate(rows):
        r = 28 + i
        assert not mp[f'B{r}'].value
        for c, v in zip('ABCDEFG', vals):
            mp[f'{c}{r}'] = v
            _st(mp[f'{c}27'], mp[f'{c}{r}'])


# ════════════════════════════════════════════════════════════════════════════
#  十、对接源_02物料：物料成本（K11）＋ 装筐费一组
# ════════════════════════════════════════════════════════════════════════════
def duijie02(wb):
    ws = wb['对接源_02物料']
    ws.merge_cells('W1:AB1')
    cell(ws, 'W1', '三、物料销售成本（跨文件取《02》【物料成本接口】A4:F27，24 个月，月末一次加权）', F_SEC, FL_SEC, AL)
    cell(ws, 'W2', '起始月', F_LBL, FL_LBL, AC)
    cell(ws, 'X2', '=IFERROR(IF([2]物料成本接口!$B$2=0,"",[2]物料成本接口!$B$2),"")', F_AUTOB, FL_AUTO, AC, fmt=MONTH)
    ws.merge_cells('Y2:AB2')
    cell(ws, 'Y2', '← 跟《02》物料成本接口 B2 一致；记账凭证自动D区按这里每月结转物料销售成本', F_NOTE, FL_NONE, AL, border=NOB)
    header(ws, 3, 'W', ['月份', '发泡网\n销售数量', '发泡网\n销售成本', '其他包装物料\n销售数量', '其他包装物料\n销售成本',
                         '销售成本合计'], 30)
    for i in range(24):
        r = 4 + i
        for j, col in enumerate(['W', 'X', 'Y', 'Z', 'AA', 'AB']):
            src = f'[2]物料成本接口!${"ABCDEF"[j]}${r}'
            cell(ws, f'{col}{r}', f'=IFERROR(IF({src}=0,"",{src}),"")', F_AUTOB if col in ('W', 'AB') else F_AUTO,
                 FL_AUTO, AC, fmt=MONTH if col == 'W' else (QTY if col in ('X', 'Z') else MONEY))
    widths(ws, {'W': 14, 'X': 11, 'Y': 13, 'Z': 12, 'AA': 13, 'AB': 13})
    # 装筐费一组（第 101～124 行，24 个月）：跟上面四组一样从往来业务明细按月汇总
    for i in range(24):
        r = 101 + i
        s = 77 + i                                      # 周转筐托盘采购入库那一组当样子
        for c in 'ABCDEFGHPQ':
            v = ws[f'{c}{s}'].value
            if isinstance(v, str) and v.startswith('='):
                v = _fill(_tmpl(v, s), r)
            ws[f'{c}{r}'] = v
            _st(ws[f'{c}{s}'], ws[f'{c}{r}'])
        ws[f'A{r}'] = r - 1
        ws[f'C{r}'] = '装筐费'
        ws[f'E{r}'] = '个'
        ws[f'D{r}'] = ws[f'D{r}'].value.replace('"筐子托盘采购"', '"装筐费"')
        ws[f'F{r}'] = ws[f'F{r}'].value.replace('"筐子托盘采购"', '"装筐费"')
        ws[f'H{r}'] = '装筐费挂应付装卸方；实际付款只从资金日记账记账'
        assert '"装筐费"' in ws[f'D{r}'].value and '"应付"' in ws[f'D{r}'].value
    # 四、建账期初物料（W29:AB91）：逐格读《02》物料成本接口隐藏网格每种物料「起始月」那一行（4、28、52…）的
    #     类别 O、月初数量 P、月初金额 Q，合计（不含筐）给【结转损益】提示【期初余额】原材料有没有登记。
    #     AB30 核对《02》网格表头没变，变了合计就留空，提示改成「见《02》其他包装物料汇总」。
    for c in ('W', 'X', 'Y', 'Z', 'AA', 'AB'):
        assert ws[f'{c}29'].value is None and ws[f'{c}{QC_TOT}'].value is None, c
    g = '[2]物料成本接口!'
    ws.merge_cells('W29:AB29')
    cell(ws, 'W29', '四、建账期初物料（取《02》【物料成本接口】每种物料起始月的月初结存，不含筐）', F_SEC, FL_SEC, AL)
    for c in ('X', 'Y', 'Z', 'AA', 'AB'):
        ws[f'{c}29'].border = BOX
    for col, t in zip(('W', 'X', 'Y', 'Z'), ('物料', '类别', '月初数量', '月初金额')):
        cell(ws, f'{col}30', t, F_HDR, FL_HDR, AC)
    cell(ws, 'AA30', '接口核对', F_LBL, FL_LBL, AC)
    cell(ws, 'AB30', (f'=IFERROR(IF(AND({g}$K$3="物料",{g}$O$3="类别",{g}$P$3="月初数量",{g}$Q$3="月初金额",'
                      f'{g}$L$5=1),"✔","✘ 接口变了"),"✘ 没取到")'), F_AUTOB, FL_AUTO, AC)
    for k in range(QC_N):
        r, s = QC0 + k, 4 + 24 * k
        cell(ws, f'W{r}', f'=IFERROR(IF({g}$K${s}="","",{g}$K${s}),"")', F_AUTO, FL_AUTO, AL)
        cell(ws, f'X{r}', f'=IF($W{r}="","",IFERROR({g}$O${s}&"",""))', F_AUTO, FL_AUTO, AC)
        cell(ws, f'Y{r}', f'=IF($W{r}="","",IFERROR(N({g}$P${s}),0))', F_AUTO, FL_AUTO, AC, fmt=QTY)
        cell(ws, f'Z{r}', f'=IF($W{r}="","",IFERROR(N({g}$Q${s}),0))', F_AUTO, FL_AUTO, AC, fmt=MONEY)
    cell(ws, f'W{QC_TOT}', '合计（不含筐）', F_TOT, FL_TOT, AC)
    cell(ws, f'X{QC_TOT}', None, F_TOT, FL_TOT, AC)
    cell(ws, f'Y{QC_TOT}', None, F_TOT, FL_TOT, AC)
    cell(ws, f'Z{QC_TOT}', (f'=IF($AB$30<>"✔","",ROUND(SUMIF($X${QC0}:$X${QC_TOT - 1},"<>筐",'
                            f'$Z${QC0}:$Z${QC_TOT - 1}),2))'), F_TOTN, FL_TOT, AC, fmt=MONEY)
    ws.merge_cells(f'AA{QC0}:AB{QC0 + 8}')
    cell(ws, f'AA{QC0}', ('← 只用来提示：【期初余额】原材料—包装物料还没登记时，【结转损益】提示栏显示这里的合计'
                          '（《02》建账期初库存＋起始月以前的进出）。AB30 显示 ✘ 时合计留空，提示改成「见《02》其他包装物料汇总」。'),
         F_NOTE, FL_NONE, AL, border=NOB)


# ════════════════════════════════════════════════════════════════════════════
#  十一、结转损益表 ＋ 记账凭证自动D区 ＋ 合计行
# ════════════════════════════════════════════════════════════════════════════
TA0 = 6            # 表 A：第 6～17 行＝12 个月
TB_H = 21          # 表 B 表头
TB0 = 22           # 表 B：第 22～51 行＝30 个损益科目
TB1 = TB0 + N_ACC - 1
TIP_COLS = {'I': 14, 'J': 14, 'K': 14, 'L': 14, 'M': 14, 'N': 14, 'O': 15}    # 表 A 提示栏 I:O 合并（也是表 B 后 6 个月＋合计列）
F_TIP = Font(name='宋体', size=9, color='006100')
# 行高按「一个月同时出 6 条提示」（现在 9 月就是这样）估：每条一行，太长的折两行
_TIP_SAMPLE = '\n'.join([
    '数据录完、核对无误后在 C 列选「是」。',
    '⚠ 本月有 87 笔「家用」记进了费用/成本，会拉低利润（要不要改走其他应收款请定）。',
    '⚠ 【原料入库计价】有 46 行有金额没填计价日期，仓储/周转/装卸费还没进总账。',
    '⚠ 物料成本含《02》建账期初库存 16,874，【期初余额】原材料—包装物料还没填，资产负债表原材料会偏低。',
    '⚠ 本月有水果销售收入 1,150.00，采购果品成本还没进总账（见【果然鲜总表】块三「采购果品」）。',
    '⚠ 本月有筐子销售收入 588.00，筐子成本还没结转：在凭证手工区补「借 其他业务成本、贷 周转材料—周转筐托盘」（金额见【果然鲜总表】块三「筐子成本」）。',
])
TIP_H = max(36, _lines(_TIP_SAMPLE, sum(TIP_COLS.values()), 9) * 12.5 + 6)


def _pz_rng(col, r1=PZ_D1):
    return f'记账凭证!${col}$4:${col}${r1}'


def build_jiezhuan(wb):
    idx = wb.sheetnames.index('利润表') + 1
    ws = wb.create_sheet(JZ, idx)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = 'C00000'
    title(ws, '结 转 损 益（按月：物料销售成本自动结转 ＋ 选「是」才生成结转本年利润凭证）', 'A', 'O')
    cell(ws, 'A2', '起始月', F_LBL, FL_LBL, AC)
    cell(ws, 'B2', datetime.datetime(2026, 8, 1), F_IN, FL_IN, AC, fmt=MONTH)
    cell(ws, 'C2', '物料成本：借', F_LBL, FL_LBL, AC)
    mp = lambda col, dflt: (f'=IF(IFERROR(INDEX(收入成本映射!${col}$4:${col}$34,MATCH("物料销售成本",收入成本映射!$B$4:$B$34,0))'
                            f'&"","")="","{dflt}",INDEX(收入成本映射!${col}$4:${col}$34,MATCH("物料销售成本",收入成本映射!$B$4:$B$34,0)))')
    cell(ws, 'D2', mp('E', '其他业务成本'), F_AUTOB, FL_AUTO, AC)
    cell(ws, 'E2', '贷', F_LBL, FL_LBL, AC)
    cell(ws, 'F2', mp('F', '原材料—包装物料'), F_AUTOB, FL_AUTO, AC)
    ws.merge_cells('G2:O2')
    cell(ws, 'G2', ('← 起始月蓝格可改（默认 2026年8月，排 12 个月）；物料成本凭证的科目在【收入成本映射】「物料销售成本」一行改。'
                    '★ 已有月份选了「是」就不能再改起始月（改了已结转的凭证会整组挪月）；满 12 个月请另起下一年账套（期初余额填本年期末）。'),
         F_NOTE, FL_NONE, AL, border=NOB)
    note(ws, '每月怎么结转：① 当月【资金日记账】和《01》《02》台账录完；② 看下面这个月的收入、成本、利润和「提示」，有 ⚠ 的先处理；'
             '③ 在 C 列把这个月改成「是」——【记账凭证】最后的自动D区就生成「记-年月-结转」凭证（借各收入类科目、贷本年利润；'
             '借本年利润、贷各成本费用类科目），损益类科目清零、利润转进本年利润；④ 资产负债表、利润表不用动（利润表自动排除结转凭证）。'
             '以后补录或修改这个月的数，结转凭证自动跟着变，不用反结账。物料销售成本那组凭证跟开关无关，有成本就自动生成。',
         'A', 'O', 3, 58, warn=True)
    section(ws, '一、每月结转开关与利润（C 列下拉：是／否，默认「否」）', 'A', 'O', 4)
    header(ws, 5, 'A', ['月份', '月末日期', '是否结转', '收入类合计', '成本费用类合计', '本月利润', '其中：物料销售成本',
                         '结转状态', '提示'], 32)
    ws.merge_cells('I5:O5')
    # 隐藏辅助：原料入库计价 有金额没填计价日期的行数
    cell(ws, 'Q2', '计价未填', F_HELP, FL_NONE, AC, border=NOB)
    cell(ws, 'R2', '=COUNTIFS(原料入库计价!$K$5:$K$904,">0")-COUNTIFS(原料入库计价!$K$5:$K$904,">0",原料入库计价!$M$5:$M$904,">0")',
         F_HELP, FL_NONE, AC, border=NOB)
    # 隐藏辅助：T2＝【期初余额】里物料成本贷方科目（F2，默认 原材料—包装物料）的期初借方净额；
    #           T3＝《02》物料成本接口起始月的月初结存合计（对接源_02物料 Z91，取不到时为空）
    cell(ws, 'S2', '期初原材料', F_HELP, FL_NONE, AC, border=NOB)
    cell(ws, 'T2', ('=ROUND(SUMIF(期初余额!$B$4:$B$81,$F$2,期初余额!$D$4:$D$81)'
                    '-SUMIF(期初余额!$B$4:$B$81,$F$2,期初余额!$E$4:$E$81),2)'), F_HELP, FL_NONE, AC, border=NOB)
    cell(ws, 'S3', '《02》期初物料', F_HELP, FL_NONE, AC, border=NOB)
    cell(ws, 'T3', f'=对接源_02物料!$Z${QC_TOT}', F_HELP, FL_NONE, AC, border=NOB)
    assert wb['期初余额']['B82'].value == '合计' and wb['期初余额']['B15'].value == '原材料—包装物料'
    jx = wb['进销存对接']
    jx1 = [r for r in range(4, jx.max_row + 1) if jx[f'D{r}'].value == '合计']
    assert len(jx1) == 1, jx1
    jx1 = jx1[0] - 1                                    # 进销存对接 数据区 4～305
    jxr = lambda c: f'进销存对接!${c}$4:${c}${jx1}'
    b2 = 'IF(N($B$2)>0,$B$2,DATE(2026,8,1))'
    for k in range(12):
        r = TA0 + k
        mc = L(3 + k)                                   # 表 B 里这个月的列（C～N）
        cell(ws, f'A{r}', f'=DATE(YEAR({b2}),MONTH({b2})+{k},1)', F_AUTOB, FL_AUTO, AC, fmt=MONTH)
        cell(ws, f'B{r}', f'=EOMONTH($A{r},0)', F_AUTO, FL_AUTO, AC, fmt=DATE)
        cell(ws, f'C{r}', '否', F_IN, FL_IN, AC)
        cell(ws, f'D{r}', f'=ROUND(SUMIF($B${TB0}:$B${TB1},"收入",{mc}${TB0}:{mc}${TB1}),2)', F_AUTO, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'E{r}', f'=ROUND(SUMIF($B${TB0}:$B${TB1},"费用",{mc}${TB0}:{mc}${TB1}),2)', F_AUTO, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'F{r}', f'=ROUND(D{r}-E{r},2)', F_AUTOB, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'G{r}', (f'=ROUND(SUMIFS(记账凭证!$H${PZ_D0}:$H${PZ_D1},记账凭证!$B${PZ_D0}:$B${PZ_D1},$B{r},'
                           f'记账凭证!$E${PZ_D0}:$E${PZ_D1},$D$2)-SUMIFS(记账凭证!$I${PZ_D0}:$I${PZ_D1},'
                           f'记账凭证!$B${PZ_D0}:$B${PZ_D1},$B{r},记账凭证!$E${PZ_D0}:$E${PZ_D1},$D$2),2)'),
             F_AUTO, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'H{r}', (f'=IF($C{r}="是",IF(AND(D{r}=0,E{r}=0),"已选是（本月无损益）","✔ 已结转"),'
                           f'IF(AND(D{r}=0,E{r}=0),"","未结转"))'), F_AUTOB, FL_AUTO, AC)
        jy = (f'COUNTIFS(资金日记账!$F$4:$F${JR1},"家用",资金日记账!$B$4:$B${JR1},">="&$A{r},资金日记账!$B$4:$B${JR1},"<="&$B{r},'
              f'资金日记账!$K$4:$K${JR1},"{{}}")')
        njy = '+'.join(jy.replace('{}', k2) for k2 in ('管理费用', '销售费用', '主营业务成本'))
        acc = lambda name: f'SUMIF($A${TB0}:$A${TB1},"{name}",{mc}${TB0}:{mc}${TB1})'     # 表 B 这个月某科目净额
        sale = acc('其他业务收入—物料销售')
        fs, bsl = acc('主营业务收入—水果销售'), acc('其他业务收入—筐子租售')
        fc = f'SUMIFS({jxr("K")},{jxr("D")},"自购水果销售成本",{jxr("C")},$A{r})'
        bc = (f'SUMIFS({_pz_rng("I")},{_pz_rng("E")},"周转材料—周转筐托盘",'
              f'{_pz_rng("B")},">="&$A{r},{_pz_rng("B")},"<="&$B{r})')                     # 本月贷记周转筐托盘
        act = f'OR(D{r}<>0,E{r}<>0)'
        nl = 'CHAR(10)&'                                  # 每条提示单独一行；最后 MID(…,2,…) 去掉开头那个换行
        parts = [
            f'IF(AND($C{r}<>"是",{act}),{nl}"数据录完、核对无误后在 C 列选「是」。","")',
            f'IF({njy}>0,{nl}"⚠ 本月有 "&({njy})&" 笔「家用」记进了费用/成本，会拉低利润（要不要改走其他应收款请定）。","")',
            (f'IF(AND(N({sale})>0,N(G{r})=0),{nl}"⚠ 本月有物料销售收入、没有物料销售成本：'
             f'看【对接源_02物料】W:AB 是否取到《02》物料成本接口。","")'),
            f'IF(AND(N($R$2)>0,{act}),{nl}"⚠ 【原料入库计价】有 "&$R$2&" 行有金额没填计价日期，仓储/周转/装卸费还没进总账。","")',
            (f'IF(AND({act},N($T$2)=0,OR(N($T$3)>0,AND($T$3="",N(G{r})<>0))),{nl}"⚠ 物料成本含《02》建账期初库存"'
             f'&IF(N($T$3)>0," "&TEXT($T$3,"#,##0"),"（金额见《02》【其他包装物料汇总】期初金额合计）")'
             f'&"，【期初余额】"&$F$2&"还没填，资产负债表原材料会偏低。","")'),
            (f'IF(AND(N({fs})>0,N({fc})=0),{nl}"⚠ 本月有水果销售收入 "&TEXT({fs},"#,##0.00")'
             f'&"，采购果品成本还没进总账（见【果然鲜总表】块三「采购果品」）。","")'),
            (f'IF(AND(N({bsl})>0,N({bc})=0),{nl}"⚠ 本月有筐子销售收入 "&TEXT({bsl},"#,##0.00")'
             f'&"，筐子成本还没结转：在凭证手工区补「借 其他业务成本、贷 周转材料—周转筐托盘」（金额见【果然鲜总表】块三「筐子成本」）。","")'),
        ]
        f_tip = '=MID(' + '&'.join(parts) + ',2,2000)'
        assert len(f_tip) < 8000, len(f_tip)
        for s in re.findall(r'"([^"]*)"', f_tip):
            assert len(s) <= 255, s
        ws.merge_cells(f'I{r}:O{r}')
        cell(ws, f'I{r}', f_tip, F_TIP, FL_AUTO, AL)
        ws.row_dimensions[r].height = TIP_H
    rt = TA0 + 12
    cell(ws, f'A{rt}', '合　计', F_TOT, FL_TOT, AC)
    for c in 'BCHI':
        cell(ws, f'{c}{rt}', None, F_TOT, FL_TOT, AC)
    cell(ws, f'C{rt}', f'=COUNTIF(C{TA0}:C{TA0 + 11},"是")&" 个月已结转"', F_TOT, FL_TOT, AC)
    for c in 'DEFG':
        cell(ws, f'{c}{rt}', f'=ROUND(SUM({c}{TA0}:{c}{TA0 + 11}),2)', F_TOTN, FL_TOT, AC, fmt=MONEY)
    ws.merge_cells(f'I{rt}:O{rt}')
    cell(ws, f'I{rt}', (f'=IF(ROUND(记账凭证!$H${PZ_TOT}-记账凭证!$I${PZ_TOT},2)=0,"记账凭证借贷平衡 ✔",'
                        f'"✘ 记账凭证借贷不平，差 "&TEXT(记账凭证!$H${PZ_TOT}-记账凭证!$I${PZ_TOT},"#,##0.00"))'),
         F_CHK, FL_TOT, AL)
    dv = DataValidation(type='list', formula1='"是,否"', allow_blank=True, showErrorMessage=True,
                        errorTitle='是否结转', error='请从下拉选「是」或「否」')
    dv.add(f'C{TA0}:C{TA0 + 11}')
    ws.add_data_validation(dv)
    # 起始月锁定：已有任一月选「是」就不让改 B2（改了已结转月份的凭证会整组挪到别的月份）
    dvd = DataValidation(type='custom',
                         formula1=f'AND(ISNUMBER($B$2),$B$2>=36526,$B$2<=73050,COUNTIF($C${TA0}:$C${TA0 + 11},"是")=0)',
                         allow_blank=False, showErrorMessage=True, errorTitle='起始月',
                         error='已有结转的月份，不能改起始月；满 12 个月请另起下一年账套（期初余额填本年期末）。'
                               '还没有结转时，请填每月 1 日的日期，如 2026/8/1。',
                         showInputMessage=True, promptTitle='起始月',
                         prompt='填每月 1 日的日期。已有月份选了「是」以后就不能再改。')
    dvd.add('B2')
    ws.add_data_validation(dvd)
    ws.conditional_formatting.add(f'I{TA0}:I{TA0 + 11}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",$I{TA0}))'],
                                                                      fill=fill('FCE4D6')))
    ws.conditional_formatting.add(f'C{TA0}:C{TA0 + 11}', FormulaRule(formula=[f'$C{TA0}="是"'], fill=fill('C6EFCE'),
                                                                      font=Font(bold=True, color='006100')))
    # 表 B：损益科目 × 12 个月
    section(ws, '二、各损益科目每月净发生额（收入类＝贷方−借方，成本费用类＝借方−贷方；已含自动D区物料销售成本，不含结转凭证本身）',
            'A', 'O', TB_H - 1)
    cell(ws, f'A{TB_H}', '科目名称', F_HDR, FL_HDR, AC)
    cell(ws, f'B{TB_H}', '类别', F_HDR, FL_HDR, AC)
    for k in range(12):
        cell(ws, f'{L(3 + k)}{TB_H}', f'=$A${TA0 + k}', F_HDR, FL_HDR, AC, fmt=MONTH)
    cell(ws, f'O{TB_H}', '12 个月合计', F_HDR, FL_HDR, AC)
    ws.row_dimensions[TB_H].height = 30
    # 隐藏辅助：科目表 4～123 行里类别＝损益的累计编号
    cell(ws, 'Q3', '损益序', F_HELP, FL_NONE, AC, border=NOB)
    for r in range(4, KM_NEW + 1):
        ws[f'Q{r}'] = f'=N(Q{r - 1})+IF(INDEX(科目表!$C$4:$C${KM_NEW},ROW()-3)="损益",1,0)'
        ws[f'Q{r}'].font = F_HELP
    for i in range(N_ACC):
        r = TB0 + i
        cell(ws, f'A{r}', (f'=IF({i + 1}>$Q${KM_NEW},"",INDEX(科目表!$B$4:$B${KM_NEW},'
                           f'MATCH({i + 1},$Q$4:$Q${KM_NEW},0)))'), F_AUTO, FL_AUTO, AL)
        cell(ws, f'B{r}', (f'=IF($A{r}="","",IF(IFERROR(INDEX(科目表!$F$4:$F${KM_NEW},MATCH($A{r},科目表!$B$4:$B${KM_NEW},0)),"")'
                           f'="收入","收入",IF(IFERROR(INDEX(科目表!$F$4:$F${KM_NEW},MATCH($A{r},科目表!$B$4:$B${KM_NEW},0)),"")="费用",'
                           f'"费用",IF(IFERROR(INDEX(科目表!$D$4:$D${KM_NEW},MATCH($A{r},科目表!$B$4:$B${KM_NEW},0)),"")="贷","收入","费用"))))'),
             F_AUTO, FL_AUTO, AC)
        for k in range(12):
            c = L(3 + k)
            crit = (f'{_pz_rng("E")},$A{r},{_pz_rng("B")},">="&{c}${TB_H},{_pz_rng("B")},"<="&EOMONTH({c}${TB_H},0)')
            cell(ws, f'{c}{r}', (f'=IF($A{r}="","",ROUND(IF($B{r}="收入",1,-1)*(SUMIFS({_pz_rng("I")},{crit})'
                                 f'-SUMIFS({_pz_rng("H")},{crit})),2))'), F_AUTO, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'O{r}', f'=IF($A{r}="","",ROUND(SUM(C{r}:N{r}),2))', F_AUTOB, FL_AUTO, AC, fmt=MONEY)
    for j, (lab, kw) in enumerate((('收入类合计', '收入'), ('成本费用类合计', '费用'))):
        r = TB1 + 1 + j
        cell(ws, f'A{r}', lab, F_TOT, FL_TOT, AC)
        cell(ws, f'B{r}', None, F_TOT, FL_TOT, AC)
        for k in range(13):
            c = L(3 + k)
            cell(ws, f'{c}{r}', f'=ROUND(SUMIF($B${TB0}:$B${TB1},"{kw}",{c}${TB0}:{c}${TB1}),2)', F_TOTN, FL_TOT, AC, fmt=MONEY)
    note(ws, f'科目清单自动取【科目表】科目类别＝损益的科目（最多 {N_ACC} 个）。本表取【记账凭证】第 4～{PZ_D1} 行'
             f'（含自动D区的物料销售成本，不含 {PZ_J0} 行以后的结转凭证本身），所以不会循环引用。'
             f'12 月（或年度最后一个月）结转后，本年利润转利润分配的分录请在手工区录（资产负债表两者都并入未分配利润，不影响平衡）。',
         'A', 'O', TB1 + 4, 44)
    widths(ws, {'A': 26, 'B': 12, 'C': 12, 'D': 13, 'E': 14, 'F': 13, 'G': 13, 'H': 14, **TIP_COLS,
                'Q': 7, 'R': 7, 'S': 7, 'T': 7})
    hide(ws, 'Q', 'R', 'S', 'T')
    # 第 2 行：G2 说明（G:O 合并）按字数给够行高；D2/F2 科目名也可能折两行
    ws.row_dimensions[2].height = max(30, _lines(ws['G2'].value, 13 + 14 + sum(TIP_COLS.values()), 9) * 12.5 + 6)
    ws.sheet_view.selection = [Selection()]
    ws.freeze_panes = 'B6'
    return ws


def _vd_common(r, src):
    return {
        'A': f'=IF($E{r}="","",ROW()-3)',
        'F': f'=IF($E{r}="","",IFERROR(INDEX(科目表!$A$4:$A${KM_NEW},MATCH($E{r},科目表!$B$4:$B${KM_NEW},0)),""))',
        'G': f'=IF($E{r}="","",IFERROR(INDEX(科目表!$E$4:$E${KM_NEW},MATCH($E{r},科目表!$B$4:$B${KM_NEW},0)),""))',
        'N': f'=IF($E{r}="","","系统")',
        'O': f'=IF($E{r}="","","{src}")',
    }


def vouchers_d(wb, tot):
    ws = wb['记账凭证']
    sty = [{c: copy.copy(ws[f'{c}{PZ_B0 + i}']._style) for c in 'ABCDEFGHIJKLMNOP'} for i in (0, 1)]
    J = JZ
    # 自动D·物料成本：24 个月 × 2 行（借 其他业务成本 / 贷 原材料—包装物料），金额＝对接源_02物料 AB
    for j in range(N_MC):
        sr = 4 + j
        amt = f'ROUND(N(对接源_02物料!$AB${sr}),2)'
        cond = f'OR(N(对接源_02物料!$W${sr})=0,{amt}=0)'
        for s in (0, 1):
            r = PZ_D0 + 2 * j + s
            v = _vd_common(r, '自动D')
            v.update({
                'B': f'=IF($E{r}="","",EOMONTH(对接源_02物料!$W${sr},0))',
                'C': f'=IF($E{r}="","","记-"&TEXT($B{r},"YYYYMM")&"-物料成本")',
                'D': f'=IF($E{r}="","","结转物料销售成本 "&TEXT($B{r},"YYYY年M月"))',
                'E': f'=IF({cond},"",{J}!${"D" if s == 0 else "F"}$2)',
                'H': f'=IF($E{r}="","",{"MAX(" + amt + ",0)" if s == 0 else "MAX(-" + amt + ",0)"})',
                'I': f'=IF($E{r}="","",{"MAX(-" + amt + ",0)" if s == 0 else "MAX(" + amt + ",0)"})',
                'K': f'=IF($E{r}="","","物料销售成本")',
            })
            for c in 'ABCDEFGHIJKLMNOP':
                d = ws[f'{c}{r}']
                d.value = v.get(c)
                d._style = sty[s][c]
    # 自动D·结转：12 个月 × 32 行
    for m in range(12):
        ra = TA0 + m                                     # 结转损益 表 A 这个月那一行
        mc = L(3 + m)
        on = f'{J}!$C${ra}<>"是"'
        base = PZ_J0 + m * BLK
        for i in range(BLK):
            r = base + i
            v = _vd_common(r, '结转')
            if i < N_ACC:
                tb = TB0 + i
                val = f'N({J}!{mc}{tb})'
                v['E'] = f'=IF({on},"",IF(OR({J}!$A{tb}="",ROUND({val},2)=0),"",{J}!$A{tb}))'
                v['H'] = f'=IF($E{r}="","",IF({J}!$B{tb}="收入",MAX({val},0),MAX(-{val},0)))'
                v['I'] = f'=IF($E{r}="","",IF({J}!$B{tb}="收入",MAX(-{val},0),MAX({val},0)))'
            else:
                col = 'D' if i == N_ACC else 'E'          # 收入类合计 → 贷本年利润；成本费用类合计 → 借本年利润
                val = f'N({J}!${col}${ra})'
                v['E'] = f'=IF(OR({on},ROUND({val},2)=0),"","本年利润")'
                if i == N_ACC:
                    v['H'] = f'=IF($E{r}="","",MAX(-{val},0))'
                    v['I'] = f'=IF($E{r}="","",MAX({val},0))'
                else:
                    v['H'] = f'=IF($E{r}="","",MAX({val},0))'
                    v['I'] = f'=IF($E{r}="","",MAX(-{val},0))'
            v['B'] = f'=IF($E{r}="","",{J}!$B${ra})'
            v['C'] = f'=IF($E{r}="","","记-"&TEXT($B{r},"YYYYMM")&"-结转")'
            v['D'] = f'=IF($E{r}="","","结转 "&TEXT($B{r},"YYYY年M月")&" 损益")'
            v['K'] = f'=IF($E{r}="","","结转损益")'
            for c in 'ABCDEFGHIJKLMNOP':
                d = ws[f'{c}{r}']
                d.value = v.get(c)
                d._style = sty[i % 2][c]
    # 合计行挪到最后
    for c, (val, s) in tot.items():
        d = ws[f'{c}{PZ_TOT}']
        if isinstance(val, str):
            val = val.replace('H4:H3609', f'H4:H{PZ_END}').replace('I4:I3609', f'I4:I{PZ_END}') \
                     .replace('H3610', f'H{PZ_TOT}').replace('I3610', f'I{PZ_TOT}')
        d.value = val
        d._style = s
    assert ws[f'H{PZ_TOT}'].value == f'=ROUND(SUM(H4:H{PZ_END}),2)', ws[f'H{PZ_TOT}'].value
    ws.auto_filter.ref = f'A3:P{PZ_TOT}'
    # A2 说明：区段行号
    a2 = ws['A2'].value
    old_tail = a2[a2.index('绿底＝自动生成勿改'):]
    old_tail = old_tail[:old_tail.index('"')]
    new_tail = (f'绿底＝自动生成勿改：A区4-909←进销存对接，B区910-2109和B区续{PZ_X0}-{PZ_X1}←资金日记账，'
                f'C区2110-3009←成本费用登记表，自动D区{PZ_D0}-{PZ_J1}＝物料销售成本和按月结转（看【结转损益】）；'
                f'黄底手工区3010-3609录调整分录。点「科目名称」筛选去掉空白，只看有内容的分录')
    assert len(new_tail) < 250
    ws['A2'] = a2.replace(old_tail, new_tail)
    for s in re.findall(r'"([^"]*)"', ws['A2'].value):
        assert len(s) <= 255, len(s)


# ════════════════════════════════════════════════════════════════════════════
#  十二、利润表、资产负债表
# ════════════════════════════════════════════════════════════════════════════
INCOME_ITEMS = {'营业收入', '投资收益', '营业外收入'}


def reports(wb):
    ws = wb['利润表']
    n = 0
    for r in range(5, 17):
        for col, d1, d2 in (('C', '$H$1', '$H$2'), ('D', '$H$3', '$H$2')):
            v = ws[f'{col}{r}'].value
            if not (isinstance(v, str) and 'SUMIFS(记账凭证!' in v):
                continue
            item = re.search(r'记账凭证!\$G\$4:\$G\$\d+,"([^"]+)"', v).group(1)
            crit = (f'记账凭证!$G$4:$G${PZ_END},"{item}",记账凭证!$O$4:$O${PZ_END},"<>结转",'
                    f'记账凭证!$B$4:$B${PZ_END},">="&{d1},记账凭证!$B$4:$B${PZ_END},"<="&{d2}')
            a, b = ('I', 'H') if item in INCOME_ITEMS else ('H', 'I')
            ws[f'{col}{r}'] = f'=ROUND(SUMIFS(记账凭证!${a}$4:${a}${PZ_END},{crit})-SUMIFS(记账凭证!${b}$4:${b}${PZ_END},{crit}),2)'
            n += 1
    assert n == 20, n
    ws['A19'] = ('提示：三张报表都按日期出数 —— 【科目余额表】管资产负债表；【利润表】可以单独设自己的区间，留空就跟着科目余额表走。'
                 '凭证没有日期的行不会被算进任何区间。利润表按净额取数（收入＝贷方−借方，费用＝借方−贷方，利息收入等贷方发生额也算进来），'
                 '并排除「结转」凭证，所以结转前后利润表的数一样。')
    ws.row_dimensions[19].height = 48
    # 资产负债表：未分配利润始终并入还没结转的损益余额（按月结转后已结转月份的损益余额是 0，不会重复）
    bs = wb['资产负债表']
    for c in ('G28', 'H28'):
        v = bs[c].value
        m = re.search(r'\+IF\(\$B\$3="否",(.*),0\),2\)$', v)
        assert m, v
        bs[c] = v[:m.start()] + '+' + m.group(1) + ',2)'
    # 第 3 行：原来的「本期是否已做结转损益凭证」下拉开关已经不起作用，改成自动显示已结转几个月（取【结转损益】合计行）
    assert bs['D3'].value is None and bs['E3'].value == '编制单位：' and bs['G3'].value == '报表日期：'
    left, hit = [], 0
    for dv in bs.data_validations.dataValidation:
        rng = [str(x) for x in dv.sqref.ranges]
        assert not any(':' in x and 'B3' in MultiCellRange(x) for x in rng), rng
        if 'B3' in rng:
            hit += 1
            rng.remove('B3')
            if not rng:
                continue
            dv.sqref = MultiCellRange(' '.join(rng))
        left.append(dv)
    assert hit == 1, hit
    bs.data_validations.dataValidation = left
    bs['A3'] = '已结转月数（看【结转损益】）：'
    bs['C3'] = None
    bs.merge_cells('B3:C3')
    cell(bs, 'B3', f'={JZ}!$C${TA0 + 12}', F_AUTOB, FL_AUTO, AC)
    cell(bs, 'C3', None, F_AUTOB, FL_AUTO, AC)        # B3:C3 合并（B 列只有 8 宽，放不下「x 个月已结转」）；原 C3 的长说明删掉，
                                                     # E3～H3 是编制单位/报表日期，没地方放，A3 已写明去【结转损益】看


# ════════════════════════════════════════════════════════════════════════════
#  十三、主页、使用说明
# ════════════════════════════════════════════════════════════════════════════
def home(wb):
    ws = wb['主页']
    r0 = ws.max_row + 2
    ws.merge_cells(f'A{r0}:D{r0}')
    ws[f'A{r0}'] = '本 轮（10/6）新 增 / 改 动 · 账 务'
    _st(ws['A42'], ws[f'A{r0}'])
    for c in 'ABCD':
        ws[f'{c}{r0 + 1}'] = ws[f'{c}43'].value
        _st(ws[f'{c}43'], ws[f'{c}{r0 + 1}'])
    items = [
        (JZ, '★每月在这里选「是」就自动生成结转凭证；物料销售成本按《02》物料成本接口每月自动结转', '录入+自动'),
        ('记账凭证', f'B区续 {PZ_X0}～{PZ_X1} 行接日记账第 604～4003 行；自动D区 {PZ_D0}～{PZ_J1} 行＝物料成本＋结转；合计在 {PZ_TOT} 行', '自动'),
        ('资金日记账', '已扩到 4,000 行：只在最后一笔下面接着录，不要在中间插行', '录入'),
        ('对账补充行', '金额直接填 H 列（蓝色），负数＝抹零/优惠冲减，自动进往来对账单', '录入'),
        ('往来业务明细', f'新段接在后面：日记账续段、周转筐续段、原料公司购买（应付）、装筐费（应付），最后一行 {WL_END}', '自动'),
    ]
    for i, (nm, txt, attr) in enumerate(items):
        r = r0 + 2 + i
        ws[f'A{r}'] = i + 1
        ws[f'B{r}'] = f'=HYPERLINK("#\'{nm}\'!A1","{nm}")'
        ws[f'C{r}'] = txt
        ws[f'D{r}'] = attr
        for c in 'ABCD':
            _st(ws[f'{c}55'], ws[f'{c}{r}'])


def manual(wb):
    ws = wb['使用说明']
    ws['B36'] = (f'　手工区 （第3010～3609行）：调整分录自己录；B区续（第{PZ_X0}～{PZ_X1}行）接日记账 604～4003 行；'
                 f'自动D区（第{PZ_D0}～{PZ_J1}行）＝物料销售成本＋按月结转（在【结转损益】选「是」）')
    r0 = ws.max_row + 2
    lines = [
        '★ 本 次 更 新（2026-10-06 · 第十版 · 账务）',
        '1、新顺发泡网、桂疆这类「包装物料采购／水果采购／周转筐采购」付款，现在一律冲应付账款（不再因为填了费用项目「发泡网」就进管理费用），'
        '往来对账单上新顺的「本期已付款」出来了。',
        '2、【对账补充行】金额直接填 H 列（蓝色）：正数＝增加欠款，负数＝冲减（抹零、优惠）。填完就进往来对账单、应收应付汇总表。'
        '注意它不进记账凭证，抹零要进总账的请在记账凭证手工区补一笔。',
        '3、【资金日记账】扩到 4,000 行（第 4～4003 行）。★只在最后一笔下面接着录，不要在中间插行——凭证、往来、各汇总都是按行位置对应的。',
        '4、每月怎么结转：当月数据录完 → 打开【结转损益】看这个月的收入、成本、利润和提示（有 ⚠ 先处理）→ C 列选「是」。'
        '记账凭证最后自动生成「记-年月-结转」凭证，损益类清零转入本年利润；资产负债表、利润表不用动。以后改了这个月的数，结转凭证自动跟着变。',
        '5、物料销售成本：按《02》【物料成本接口】（月末一次加权平均）每月自动结转「借 其他业务成本 贷 原材料—包装物料」，跟结转开关无关。',
        '6、【利润表】改成按净额取数并排除结转凭证（利息收入不再漏掉）；【资产负债表】B3 改成自动显示「几个月已结转」（取【结转损益】），不用再选。',
        '7、【往来业务明细】新接四段：日记账 604 行以后、周转筐 2005 行以后、《01》原料出库「公司购买」（应付卖方）、《02》装筐费（应付装卸方）。',
        '8、【结转损益】起始月（B2）只能在还没有任何月份选「是」的时候改，选过「是」就锁住（改了已结转的凭证会整组挪到别的月份）。'
        '满 12 个月请另起下一年账套：复制一份，把本年期末余额填进新账套的【期初余额】。',
        '9、【结转损益】提示栏新增三条：【期初余额】原材料—包装物料还没登记《02》建账期初物料（金额取《02》物料成本接口）；'
        '本月有水果销售收入、采购果品成本还没进总账；本月卖了筐子、筐子成本还没结转（在凭证手工区补「借 其他业务成本 贷 周转材料—周转筐托盘」）。',
    ]
    width = sum(ws.column_dimensions[c].width if c in ws.column_dimensions else ws.sheet_format.defaultColWidth or 9
                for c in 'ABCDEFGH')
    for i, t in enumerate(lines):
        r = r0 + i
        ws.merge_cells(f'A{r}:H{r}')
        ws[f'A{r}'] = t
        _st(ws['A123' if i == 0 else 'A124'], ws[f'A{r}'])
        h0 = ws.row_dimensions[123 if i == 0 else 124].height or (20 if i == 0 else 36)
        ws.row_dimensions[r].height = h0 if i == 0 else max(h0, _lines(t, width, 10) * 15 + 4)


# ════════════════════════════════════════════════════════════════════════════
def apply(wb):
    assert 'B区910-2109' in wb['记账凭证']['A2'].value or '910-2109' in wb['记账凭证']['A2'].value
    widen_all(wb)            # 先放宽现有公式的写死范围（之后新写的公式直接用新范围）
    journal(wb)
    helpers(wb)
    monthly_rows(wb)
    pool(wb)
    tot = vouchers_b(wb)
    wanglai(wb)
    buchong(wb)
    kemu(wb)
    duijie02(wb)
    build_jiezhuan(wb)
    vouchers_d(wb, tot)
    reports(wb)
    home(wb)
    manual(wb)


def post(path):
    """externalLink 里 WPS 写的 02 绝对路径改回相对路径"""
    tmp = path + '.tmp03a'
    zin = zipfile.ZipFile(path)
    zout = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    n = 0
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename.startswith('xl/externalLinks/_rels/') and it.filename.endswith('.rels'):
            s = data.decode('utf8')
            s2 = re.sub(r'Target="[^"]*/(0[1-4]_[^"/]+\.xlsx)"', r'Target="\1"', s)
            if s2 != s:
                n += 1
                data = s2.encode('utf8')
        zout.writestr(it, data)
    zout.close()
    zin.close()
    shutil.move(tmp, path)
    return n
