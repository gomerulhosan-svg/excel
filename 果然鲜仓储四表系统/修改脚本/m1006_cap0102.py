# -*- coding: utf-8 -*-
"""A050 容量（标签 cap0102）：《01》《02》按一年用量补齐容量，主页加「容量一览」。

用户原话：「A050科目表里费用项目我增加到124行添加的就无效需要调到200行，还有资金日记账的计算行数就到603，
这些都要预留到一年用量，包括其他的地方检查一下也一样。」——这个模块管《01》《02》两本。

集成（build_1006.py）：
   《01》[(m1006_01, 'apply'), (m1006_cap0102, 'apply_01'), (m1006_99, 'apply_01')]
   《02》[(m1006_02, 'apply'), (m1006_cap0102, 'apply_02')]
   都是在 openpyxl 工作簿上就地改（load_workbook(底稿)，不用 data_only），不需要 post()。
   必须跑在 m1006_01 / m1006_02 之后（要用到它们扩好的行号：原料入库 5003、对账明细接口 6003、周转筐 3003、
   装筐费接口 503 等）；对同一份底稿重复跑结果一样。

《01》apply_01：
 1. 【实盘库存盘点】账面区 5～304（300 个组合）→ 5～404（400 个，跟【库存结余】【库存等级接口】的 400 个
    货主＋品种＋等级组合一致）；合计行 305 → 405，合计/条件格式跟着改。
 2. 【基础资料】A～P 各下拉列表每列 100 项（4～103）→ 200 项（4～203）：命名区（品种表、等级表、货主表…）的
    COUNTA 区域、格线、标题跟着改；_自动清单 BQ～BS「单位」池 4～103 → 4～203。（等级一列已有 48 项，还在涨。）
 3. 【主页】加「五、容量一览」：每张明细表、接口、自动清单列表 已用/容量/占比，占比≥80% 整行变红；
    写明明细表新数据接在最后一行下面录。
 4. 【使用说明】末尾补一条。
《02》apply_02：
 1. 【装筐费接口】4～503（500 笔）→ 4～1503（1,500 笔），接口列 A～I 含义不变；A2 超量提示、条件格式改成 >1500。
    （_自动清单 ED/EE「装筐费>0/累计」池子本来就覆盖周转筐 4～3003 全部行，不用动。）
    《03》往来业务明细「02装筐费」段要按 1,500 行读（[2]装筐费接口!$x$4:$x$1503）——由《03》实施者改。
 2. 【基础资料】物料价目表 N～P（物料名称/规格/计量单位）只到第 34 行（31 种，已用 22 种），
    【包装物料出入库明细】F/G 列按它带规格、单位——第 35 行以后加的物料带不出来（跟科目表 124 行同一个病）。
    → 扩到第 103 行（100 种，跟下拉命名区 物料名称表/规格表/计量单位表 一样），格线跟着画；
    _自动清单 CE/CF「单位/物料次数」池 4～53 → 4～103。
 3. 【主页】加「四、容量一览」（同《01》）；【使用说明】末尾补一条。
"""
import os
import re
import sys
import copy

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font, PatternFill, Alignment

from m1006_02 import map_workbook, map_refs, row_end_extender, tr, shift_area, _split, _join

AU = '_自动清单'
FL_RED = PatternFill('solid', fgColor='FFC7CE')
F_RED = Font(name='微软雅黑', size=10, bold=True, color='9C0006')
RED_AT = 0.8                      # 占比≥80% 标红


def _copy_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)


def _extend_rows(ws, src_row, r0, r1, cols=None):
    """把 src_row 那一行的公式（相对行号平移）和样式复制到 r0～r1 行；没有公式的格只带样式、值清空"""
    cols = cols or range(1, ws.max_column + 1)
    h = ws.row_dimensions[src_row].height
    for c in cols:
        src = ws.cell(src_row, c)
        f = src.value if isinstance(src.value, str) and src.value.startswith('=') else None
        for r in range(r0, r1 + 1):
            dst = ws.cell(r, c)
            dst.value = tr(f, r - src_row) if f else None
            _copy_style(src, dst)
    if h:
        for r in range(r0, r1 + 1):
            ws.row_dimensions[r].height = h


def _cf_extend(ws, fn):
    """条件格式的范围按 fn 改写（规则不动）"""
    new = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        rng = shift_area(str(cf.sqref), fn, ws.title)
        for rule in cf.rules:
            new.add(rng, rule)
    ws.conditional_formatting = new


def _col_ext(sheet, cols, start, old_end, new_end):
    """只加长指定列的「start:old_end」区域（绝对结束行）"""
    def fn(sh, ref):
        if sh != sheet or ':' not in ref:
            return ref
        ps = _split(ref)
        if (len(ps) == 2 and ps[0][1] in cols and ps[1][1] in cols and ps[0][3] and ps[1][3]
                and int(ps[0][3]) == start and int(ps[1][3]) == old_end and ps[1][2] == '$'):
            ps[1][3] = str(new_end)
        return _join(ps)
    return fn


def _w(t):
    """显示宽度（汉字算 1，半角算 0.5），向上取整"""
    return int(sum(1 if ord(ch) > 0x2E7F else 0.5 for ch in str(t)) + 0.999)


def _append_manual(wb, text):
    """【使用说明】最后一条更新记录下面接一条（样式照上一条）"""
    ws = wb['使用说明']
    last = max(r for (r, c), x in ws._cells.items() if x.value not in (None, ''))
    if any(isinstance(x.value, str) and text[:20] in x.value for x in ws._cells.values()):
        return
    src = next(ws.cell(last, c) for c in range(1, ws.max_column + 1) if ws.cell(last, c).value not in (None, ''))
    m = re.match(r'(\d+)、', str(src.value))
    n = int(m.group(1)) + 1 if m else 1
    dst = ws.cell(last + 1, src.column)
    dst.value = f'{n}、{text}'
    _copy_style(src, dst)
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == last and mr.max_row == last and mr.min_col == src.column:
            ws.merge_cells(start_row=last + 1, end_row=last + 1, start_column=mr.min_col, end_column=mr.max_col)
    ws.row_dimensions[last + 1].height = max(ws.row_dimensions[last].height or 0, 36)


# ════════════════════════════════════════════════════════════════════════════
#  主页「容量一览」
# ════════════════════════════════════════════════════════════════════════════
def overview(wb, start, title, items, sec_ref, hdr_ref, row_refs):
    """主页从 start 行起写「容量一览」块。
       items: [(表名(超链接目标), 显示名, 已用公式(不带=), 容量, 单位, 口径说明)]
       sec_ref / hdr_ref / row_refs：照抄样式的格（分区标题、表头、数据行 A～D）"""
    ws = wb['主页']
    st_sec = copy.copy(ws[sec_ref]._style)
    st_hdr = {c: copy.copy(ws[f'{c}{hdr_ref}']._style) for c in 'ABCD'}
    st_row = {c: copy.copy(ws[f'{c}{row_refs}']._style) for c in 'ABCD'}
    h_sec = ws.row_dimensions[int(re.sub(r'\D', '', sec_ref))].height or 24
    r = start
    ws.merge_cells(f'A{r}:D{r}')
    ws[f'A{r}'] = title
    ws[f'A{r}']._style = st_sec
    ws.row_dimensions[r].height = h_sec
    r += 1
    ws.merge_cells(f'A{r}:D{r}')
    ws[f'A{r}'] = ('★ 明细表的新数据一律接在最后一行下面录，不要在中间插行、删行（插行会让自动清单错位、漏笔）。'
                   '下面「已用/容量」到 80% 整行变红：请联系维护人员用脚本扩容（几本表要一起扩，自己往下拖公式不够）。')
    ws[f'A{r}']._style = copy.copy(st_row['C'])
    ws[f'A{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws[f'A{r}'].font = Font(name='微软雅黑', size=10, bold=True, color='C00000')
    ws.row_dimensions[r].height = 36
    r += 1
    for c, t in zip('ABCD', ['序号', '表 / 清单', '已用 / 容量　（口径）', '占比']):
        ws[f'{c}{r}'] = t
        ws[f'{c}{r}']._style = copy.copy(st_hdr[c])
    ws.row_dimensions[r].height = ws.row_dimensions[hdr_ref].height or 19.5
    r += 1
    r0 = r
    for i, (sheet, name, used, cap, unit, desc) in enumerate(items, 1):
        ws[f'A{r}'] = i
        ws[f'B{r}'] = name
        if sheet:
            ws[f'B{r}'].hyperlink = Hyperlink(ref=f'B{r}', location=f"'{sheet}'!A1", display=name)
        txt = f'=TEXT({used},"#,##0")&" / "&TEXT({cap},"#,##0")&" {unit}　{desc}"'
        assert len(desc) + len(unit) < 200, desc
        ws[f'C{r}'] = txt
        ws[f'D{r}'] = f'=IFERROR(({used})/{cap},0)'
        for c in 'ABCD':
            ws[f'{c}{r}']._style = copy.copy(st_row[c])
        ws[f'C{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        ws[f'D{r}'].number_format = '0%'
        ws[f'D{r}'].alignment = Alignment(horizontal='center', vertical='center')
        ws[f'D{r}'].font = Font(name='微软雅黑', size=10, bold=True, color='006100')
        ws[f'D{r}'].fill = PatternFill('solid', fgColor='E2EFDA')
        ws[f'B{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        lines = max(-(-_w(name) // 11), -(-(_w(desc + unit) + 8) // 30), 1)
        ws.row_dimensions[r].height = {1: 18.75, 2: 32}.get(lines, 15 * lines)
        r += 1
    r1 = r - 1
    ws.conditional_formatting.add(f'A{r0}:D{r1}', FormulaRule(formula=[f'$D{r0}>={RED_AT}'], fill=FL_RED, font=F_RED))
    if ws.print_area:
        ws.print_area = f'A1:D{r1}'
    return r0, r1


# ════════════════════════════════════════════════════════════════════════════
#  《01》
# ════════════════════════════════════════════════════════════════════════════
PD = '实盘库存盘点'
PD_OLD1, PD1 = 304, 404           # 账面区末行 304 → 404；合计行 305 → 405


def pandian_400(wb):
    ws = wb[PD]
    assert ws[f'B{PD_OLD1 + 1}'].value == '合计', ws[f'B{PD_OLD1 + 1}'].value
    ncol = ws.max_column
    tot = {c: (ws.cell(PD_OLD1 + 1, c).value, copy.copy(ws.cell(PD_OLD1 + 1, c)._style)) for c in range(1, ncol + 1)}
    tot_h = ws.row_dimensions[PD_OLD1 + 1].height
    _extend_rows(ws, PD_OLD1, PD_OLD1 + 1, PD1, range(1, ncol + 1))
    ext = row_end_extender(PD, 5, PD_OLD1, PD1, absolute_only=False)
    for c, (v, st) in tot.items():
        d = ws.cell(PD1 + 1, c)
        if isinstance(v, str) and v.startswith('='):
            v = map_refs(tr(v, PD1 - PD_OLD1), PD, ext)
        d.value = v
        d._style = st
    if tot_h:
        ws.row_dimensions[PD1 + 1].height = tot_h
    _cf_extend(ws, ext)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(shift_area(str(dv.sqref), ext, PD))
    if ws.print_area:
        ws.print_area = [shift_area(a, ext, PD).split('!')[-1] for a in str(ws.print_area).split(',')]
    map_workbook(wb, row_end_extender(PD, 5, PD_OLD1, PD1), only_if=PD)


JC1_OLD1, JC1_1 = 103, 203         # 《01》基础资料 每列 100 项 → 200 项


def jc01_200(wb):
    """《01》基础资料 A～P 各下拉列表 4～103（每列 100 项）→ 4～203（200 项）：
       命名区 品种表/等级表/货主表…的 COUNTA 区域、格线、标题，以及 _自动清单 BQ～BS「单位」池跟着加长。
       （等级一列 10/6 已有 48 项，9/28 是 45 项，按品种×等级还会涨，留足一年。）"""
    ws = wb['基础资料']
    for c in range(1, 17):
        src = ws.cell(JC1_OLD1, c)
        for r in range(JC1_OLD1 + 1, JC1_1 + 1):
            d = ws.cell(r, c)
            assert d.value in (None, ''), (c, r, d.value)
            _copy_style(src, d)
    h = ws.row_dimensions[JC1_OLD1].height
    for r in range(JC1_OLD1 + 1, JC1_1 + 1):
        ws.row_dimensions[r].height = h
    if isinstance(ws['A1'].value, str):
        ws['A1'] = ws['A1'].value.replace('每列最多 100 项', '每列最多 200 项')
    n = 0
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text and '基础资料!$A$4:$P$103' in dn.attr_text:
            dn.attr_text = dn.attr_text.replace('基础资料!$A$4:$P$103', '基础资料!$A$4:$P$203')
            n += 1
    assert n >= 16, n
    au = wb[AU]
    for c in ('BQ', 'BR', 'BS'):
        src = au[f'{c}{JC1_OLD1}']
        for r in range(JC1_OLD1 + 1, JC1_1 + 1):
            assert au[f'{c}{r}'].value is None, (c, r)
            au[f'{c}{r}'] = tr(src.value, r - JC1_OLD1)
            _copy_style(src, au[f'{c}{r}'])
    fn = _col_ext(AU, ('BQ', 'BR', 'BS'), 4, JC1_OLD1, JC1_1)
    for ref in ('BQ1', 'BR1'):
        au[ref] = map_refs(au[ref].value, AU, fn)


def home_01(wb):
    ws = wb['主页']
    last = max(r for (r, c), x in ws._cells.items() if x.value not in (None, ''))
    jc = ','.join(f'COUNTA(基础资料!${c}$4:${c}${JC1_1})' for c in 'ABCDEFGHIJKLMNOP')
    items = [
        ('原料入库明细', '原料入库明细', 'MAX(原料入库明细!$A$4:$A$5003)', 5000, '行', '按 A 列序号（录到第几笔）'),
        ('原料出库明细', '原料出库明细', 'MAX(原料出库明细!$A$4:$A$3004)', 3001, '行', '按 A 列序号'),
        ('成品出库明细', '成品出库明细', 'MAX(成品出库明细!$A$4:$A$3003)', 3000, '行', '按 A 列序号'),
        ('其他代存明细', '其他代存明细', 'MAX(其他代存明细!$A$4:$A$803)', 800, '行', '按 A 列序号'),
        ('对账明细接口', '对账明细接口', f'N({AU}!$BX$1)', 7000, '笔', '四张明细合计笔数，《04》水果对账单读这张'),
        ('库存结余', '库存结余 / 实盘库存盘点 / 库存等级接口', f'N({AU}!$A$1)', 400, '个',
         '货主＋品种＋等级组合（含只在成品出库出现的组合）'),
        ('库存总结余', '库存总结余 / 入库吨位汇总', f'N({AU}!$S$1)', 300, '个', '货主＋品种组合（按原料入库）'),
        ('成品出库吨位汇总', '成品出库吨位汇总', f'N({AU}!$AC$1)', 300, '个', '货主＋品种组合（按成品出库）'),
        ('其他代存汇总', '其他代存汇总', f'N({AU}!$AM$1)', 200, '个', '代存客户＋品种＋等级组合'),
        ('筐子进出接口', '筐子进出接口', f'N({AU}!$BA$1)', 300, '个', '客户＋筐子类型组合，《02》对接源读这张'),
        ('果然鲜采购明细', '往来单位（财务取数接口 / 果然鲜采购明细·按货主）', f'N({AU}!$BK$1)', 120, '家',
         '财务取数接口能放 150 家、果然鲜采购明细按货主汇总 120 家，按小的算'),
        ('果然鲜采购明细', '果然鲜采购明细（逐笔）', 'MAX(果然鲜采购明细!$A$4:$A$403)', 400, '笔', '按当前日期范围；不填起止日期＝全部'),
        ('果然鲜销售明细', '果然鲜销售明细（逐笔）', 'MAX(果然鲜销售明细!$A$4:$A$403)', 400, '笔', '按当前日期范围；不填起止日期＝全部'),
        ('果然鲜销售明细', '果然鲜销售明细·按购买方', f'N({AU}!$CD$1)', 120, '家', '第 480 行起的购买方汇总'),
        ('公司购买接口', '公司购买接口', 'N(公司购买接口!$B$1)', 400, '笔', '原料出库「公司购买」，《03》挂应付读这张'),
        ('基础资料', '基础资料（下拉选项）', f'MAX({jc})', 200, '项', '每列最多 200 项（第 4～203 行），按最多的一列算'),
    ]
    overview(wb, last + 2, '五、容量一览（已用 / 容量 · 到 80% 变红）', items, 'A27', 28, 29)


def apply_01(wb):
    pandian_400(wb)
    jc01_200(wb)
    home_01(wb)
    _append_manual(wb, '【主页】加了「容量一览」：每张明细表、接口和自动列表已用多少、能放多少，到 80% 整行变红。'
                       '【实盘库存盘点】账面区扩到 400 个组合（第 5～404 行，合计在第 405 行），跟【库存结余】一致；'
                       '【基础资料】每列能填 200 项（第 4～203 行），下拉自动带出。'
                       '新数据一律接在最后一行下面录，不要在中间插行。')


# ════════════════════════════════════════════════════════════════════════════
#  《02》
# ════════════════════════════════════════════════════════════════════════════
ZKF = '装筐费接口'
ZKF_OLD1, ZKF1 = 503, 1503
JC = '基础资料'
JG_OLD1, JG1 = 34, 103            # 物料价目表 N～P 末行
CE_OLD1, CE1 = 53, 103            # _自动清单 CE/CF 单位池末行


def fee_iface_1500(wb):
    ws = wb[ZKF]
    assert ws[f'H{ZKF_OLD1}'].value and ws[f'A{ZKF_OLD1 + 1}'].value is None
    _extend_rows(ws, ZKF_OLD1, ZKF_OLD1 + 1, ZKF1, range(1, 10))
    msg = ('★ 按【周转筐出入库明细】O 列「装筐费」逐笔列出金额>0 的行（第 4 行起连续排列，最多 1,500 笔）。'
           '付款对象＝该行「装卸」列。装筐费单价在【基础资料】U～V 列设，默认空＝不计费，所以本表平时是空的。')
    over = ('★★ 注意：装筐费已有 "&N($L$1)&" 笔，超过本表 1,500 笔上限，第 1,501 笔起《03》不会挂应付，'
            '请联系维护人员扩容（《03》往来业务明细要一起扩）')
    assert len(msg) <= 255 and len(over) <= 255
    ws['A2'] = f'=IF(N($L$1)>{ZKF1 - 3},"{over}","{msg}")'
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.formula:
                rule.formula = [x.replace('N($L$1)>500', f'N($L$1)>{ZKF1 - 3}') for x in rule.formula]
    home = wb['主页']
    for c in home._cells.values():
        if c.value == '【新】装筐费逐笔清单，给《03》挂应付':
            c.value = '【新】装筐费逐笔清单（1,500 行），给《03》挂应付'


def price_list_100(wb):
    """基础资料 物料价目表 N～P 4～34 → 4～103；引用它的公式、_自动清单 单位池跟着加长"""
    ws = wb[JC]
    assert ws['N3'].value == '物料名称' and ws['P3'].value == '计量单位'
    for c in 'NOP':
        src = ws[f'{c}{JG_OLD1}']
        for r in range(JG_OLD1 + 1, JG1 + 1):
            d = ws[f'{c}{r}']
            assert d.value in (None, ''), (c, r, d.value)
            _copy_style(src, d)
    map_workbook(wb, _col_ext(JC, ('N', 'O', 'P'), 4, JG_OLD1, JG1), only_if=JC)
    map_workbook(wb, _col_ext(JC, ('P',), 4, CE_OLD1, JG1), only_if=JC)
    au = wb[AU]
    assert au[f'CE{CE_OLD1 + 1}'].value is None and au[f'CF{CE_OLD1 + 1}'].value is None
    for c in ('CE', 'CF'):
        src = au[f'{c}{CE_OLD1}']
        for r in range(CE_OLD1 + 1, CE1 + 1):
            au[f'{c}{r}'] = tr(src.value, r - CE_OLD1)
            _copy_style(src, au[f'{c}{r}'])
    au['CE1'] = map_refs(au['CE1'].value, AU, _col_ext(AU, ('CE', 'CF'), 4, CE_OLD1, CE1))


def home_02(wb):
    ws = wb['主页']
    last = max(r for (r, c), x in ws._cells.items() if x.value not in (None, ''))
    jc = ','.join(f'COUNTA(基础资料!${c}$4:${c}$103)' for c in 'ABCDEFGHIJKLNOP')
    items = [
        ('包装物料出入库明细', '包装物料出入库明细', 'MAX(包装物料出入库明细!$A$4:$A$2003)', 2000, '行', '按 A 列序号（录到第几笔）'),
        ('周转筐出入库明细', '周转筐出入库明细', 'MAX(周转筐出入库明细!$A$4:$A$3003)', 3000, '行', '按 A 列序号'),
        ('客户自备加工筐明细', '客户自备加工筐明细', 'MAX(客户自备加工筐明细!$A$4:$A$1204)', 1201, '行', '按 A 列序号'),
        ('次果筐+托盘明细', '次果筐+托盘明细', "MAX('次果筐+托盘明细'!$A$4:$A$603)", 600, '行', '按 A 列序号'),
        ('对账明细接口', '对账明细接口', f'N({AU}!$CP$1)', 4000, '笔', '四张明细合计笔数，《04》物料/筐子对账单读这张'),
        ('装筐费接口', '装筐费接口', f'N({AU}!$ED$1)', 1500, '笔', '装筐费>0 的周转筐行，《03》挂应付读这张'),
        ('物料库存结余', '物料种类（物料库存结余）', f'N({AU}!$A$1)', 50, '种', '包装物料出入库明细里出现过的物料种类'),
        ('基础资料', '物料价目表（基础资料 N～P 列）', 'COUNTA(基础资料!$N$4:$N$103)', 100, '种', '第 4～103 行，规格、单位按这里自动带'),
        ('筐子库存汇总', '筐子类型（筐子库存汇总）', f'N({AU}!$BM$1)', 60, '种', '周转筐明细＋《01》带来的全部筐型'),
        ('筐子往来汇总', '筐子往来汇总', f'N({AU}!$BY$1)', 250, '家', '周转筐对方＋《01》筐客户＋次果筐客户'),
        ('筐子销售汇总', '筐子销售汇总', f'N({AU}!$W$1)', 200, '家', '周转筐明细的对方单位'),
        ('筐子采购汇总', '筐子采购汇总', f'N({AU}!$CS$1)', 200, '家', '周转筐采购入库的供应商'),
        ('自备筐汇总', '自备筐汇总', f'N({AU}!$AI$1)', 200, '家', '自备筐明细＋《01》带自备筐的客户'),
        ('次果筐汇总', '次果筐汇总', f'N({AU}!$AU$1)', 200, '家', '次果筐/托盘客户（筐子类型最多 4 种）'),
        ('押金汇总', '押金汇总', f'N({AU}!$BS$1)', 200, '家', '往来单位'),
        ('客户领用物料汇总', '客户领用物料汇总', f'N({AU}!$DF$1)', 200, '家', '有领用/销售/退回的客户'),
        ('发泡网销售汇总', '发泡网销售汇总', f'N({AU}!$DL$1)', 200, '家', '买过发泡网的客户'),
        ('其他包装物料汇总', '其他包装物料汇总·客户', f'N({AU}!$DR$1)', 200, '家', '买过其他包装物料的客户'),
        ('其他包装物料汇总', '其他包装物料汇总·物料', f'N({AU}!$EG$1)', 20, '种', '发泡网、筐子以外的物料种类'),
        ('包装物料采购汇总', '包装物料采购汇总', f'N({AU}!$DX$1)', 200, '家', '包装物料采购入库的供应商'),
        ('财务取数接口', '财务取数接口·按客户往来', f'COUNTIF({AU}!$CL$4:$CL$203,"?*")', 150, '家', '基础资料 客户/领取人＋供应商'),
        ('对接源_01水果筐数', '对接源_01水果筐数', f'N({AU}!$BA$1)', 300, '个', '《01》筐子进出接口的客户＋筐型组合'),
        ('基础资料', '基础资料（下拉选项）', f'MAX({jc})', 100, '项', '每列最多 100 项（第 4～103 行），按最多的一列算'),
    ]
    overview(wb, last + 2, '四、容量一览（已用 / 容量 · 到 80% 变红）', items, 'A29', 30, 31)


def apply_02(wb):
    fee_iface_1500(wb)
    price_list_100(wb)
    home_02(wb)
    _append_manual(wb, '【主页】加了「容量一览」：每张明细表、接口和自动列表已用多少、能放多少，到 80% 整行变红。'
                       '【装筐费接口】扩到 1,500 笔；【基础资料】物料价目表（N～P 列）能填到第 103 行'
                       '（原来只到第 34 行，再往下加的物料带不出规格和单位）。新数据一律接在最后一行下面录，不要在中间插行。')
