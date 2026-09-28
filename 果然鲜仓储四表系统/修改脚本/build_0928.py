# -*- coding: utf-8 -*-
"""0928 这一轮：从你 9 月 28 日发来的四本（存档在 ../参考/0928上传/）改出成品。

   《01》《02》这一轮不用改，原样带过去；《03》《04》用 openpyxl 改，
   再在「副本」上用 LibreOffice 整册重算，只把算出来的数写回成品的公式格子里
   （成品本身不经 LibreOffice 存盘，字体、格式、批注、下拉都保持原样）。

跑法：python3 build_0928.py [输出目录]
"""
import os, sys, shutil, subprocess, zipfile, json, glob
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
SRC = os.path.join(HERE, '..', '参考', '0928上传')
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..')
CALC = os.path.join(OUT, '_calc')
NAMES = {'01': '01_水果进销存台账模板.xlsx', '02': '02_物料与周转物台账模板.xlsx',
         '03': '03_财务账套与报表模板.xlsx', '04': '04_综合查询对账单模板.xlsx'}
RECALC = (glob.glob('/mnt/skills/public/xlsx/scripts/recalc.py') +
          glob.glob('/root/.claude/skills/synced/*/xlsx/scripts/recalc.py'))[0]

from inject_cache import inject
import m0928_往来, m0928_资金汇总, m0928_主页与04 as m3

src = lambda k: os.path.join(SRC, NAMES[k])
out = lambda k: os.path.join(OUT, NAMES[k])


def check_links(path):
    """新版 openpyxl 存盘会保留跨文件链接（老版本会丢，才有 repair_ext_links.py）。这里只核对还在。"""
    z = zipfile.ZipFile(path)
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf8')
    n = rels.count('relationships/externalLink"')
    parts = [x for x in z.namelist() if x.startswith('xl/externalLinks/externalLink')]
    z.close()
    if n != 2 or len(parts) != 2:
        raise SystemExit(f'跨文件链接丢了：{path} 关系 {n} 条，部件 {parts}')
    print(f'   跨文件链接 2 条都在')


def hidden_in_merges(path):
    """找出藏在合并格里（不是左上角那格）却有值/公式的格子 —— openpyxl 会把它们丢掉"""
    from lxml import etree
    from openpyxl.utils.cell import range_boundaries, get_column_letter as CL
    NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    RNS = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
    z = zipfile.ZipFile(path)
    wbx = etree.fromstring(z.read('xl/workbook.xml'))
    rid = {r.get('Id'): r.get('Target') for r in etree.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    out = set()
    for sh in wbx.iter(NS + 'sheet'):
        t = rid[sh.get(RNS + 'id')].lstrip('/')
        t = t if t.startswith('xl/') else 'xl/' + t
        root = etree.fromstring(z.read(t))
        has = {c.get('r') for c in root.iter(NS + 'c')
               if c.find(NS + 'f') is not None or (c.find(NS + 'v') is not None and (c.find(NS + 'v').text or '') != '')}
        for m in root.iter(NS + 'mergeCell'):
            c1, r1, c2, r2 = range_boundaries(m.get('ref'))
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    ref = f'{CL(c)}{r}'
                    if (r, c) != (r1, c1) and ref in has:
                        out.add((sh.get('name'), ref))
    return out


def cell_present(path, sheet, ref):
    wb = openpyxl.load_workbook(path, read_only=True)
    v = wb[sheet][ref].value
    wb.close()
    return v not in (None, '')


def fonts(path):
    """前面几张表的字体被早先一次 LibreOffice 存盘换成了 Linux 字体，统一改回微软雅黑"""
    tmp = path + '.tmp'
    zi = zipfile.ZipFile(path)
    zo = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in zi.infolist():
        d = zi.read(it.filename)
        if it.filename == 'xl/styles.xml':
            s = d.decode('utf8')
            for f in ('WenQuanYi Zen Hei', 'Noto Sans CJK SC'):
                s = s.replace(f'val="{f}"', 'val="微软雅黑"')
            d = s.encode('utf8')
        zo.writestr(it, d)
    zo.close(); zi.close()
    shutil.move(tmp, path)


def recalc(path, timeout=1800):
    r = subprocess.run([sys.executable, RECALC, path, str(timeout), '--force'], capture_output=True, text=True)
    try:
        j = json.loads(r.stdout[r.stdout.index('{'):])
    except Exception:
        print(r.stdout, r.stderr); raise SystemExit('重算失败')
    print('   重算', os.path.basename(path), j.get('status'), '公式', j.get('total_formulas'), '错误', j.get('total_errors'))
    if j.get('total_errors'):
        print('   ', json.dumps(j.get('error_summary'), ensure_ascii=False)[:1500])
    return j


if __name__ == '__main__':
    os.makedirs(OUT, exist_ok=True)
    print('① 《01》《02》原样带过去')
    for k in ('01', '02'):
        shutil.copy(src(k), out(k))
    print('② 改《03》')
    wb = openpyxl.load_workbook(src('03'))
    m0928_往来.apply(wb)
    m0928_资金汇总.apply(wb)
    m3.home_03(wb)
    m3.manual_03(wb)
    m3.finish_03(wb)
    wb.save(out('03'))
    check_links(out('03'))
    for sh, ref in sorted(hidden_in_merges(src('03'))):
        if not cell_present(out('03'), sh, ref):
            raise SystemExit(f'合并格里藏的值丢了：{sh}!{ref}')
        print(f'   合并格里藏的值已保住：{sh}!{ref}')
    fonts(out('03'))
    print('③ 改《04》')
    wb = openpyxl.load_workbook(src('04'))
    m3.buyer_04(wb)
    wb.save(out('04'))
    check_links(out('04'))
    for sh, ref in sorted(hidden_in_merges(src('04'))):
        if not cell_present(out('04'), sh, ref):
            raise SystemExit(f'合并格里藏的值丢了：{sh}!{ref}')
    fonts(out('04'))
    print('④ 在副本上整册重算，把数写回成品')
    os.makedirs(CALC, exist_ok=True)
    for k in NAMES:
        shutil.copy(out(k), os.path.join(CALC, NAMES[k]))
    res = {}
    for k in ('03', '04'):
        res[k] = recalc(os.path.join(CALC, NAMES[k]))
        n, s = inject(out(k), os.path.join(CALC, NAMES[k]))
        print(f'   {NAMES[k]} 写回缓存 {n} 格（空 {s} 格）')
    if '--keep' not in sys.argv:
        shutil.rmtree(CALC, ignore_errors=True)
    print('\n完成，成品在', os.path.abspath(OUT))
