# -*- coding: utf-8 -*-
"""0928 这一轮新做/重做的表共用的东西：前面几张老表的样式（深蓝抬头、蓝表头、浅绿公式格、
浅蓝录入格、细灰格线），以及「年份 / 月份 / 起始日 / 截止日」时段选择块。"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI

# ── 样式：照【资金日记账】【科目余额表】【收入月度汇总】这些前面的表 ─────────────
HEI = '微软雅黑'
SONG = '宋体'
F_TITLE = Font(name=HEI, size=14, bold=True, color='FFFFFF')
F_NOTE = Font(name=HEI, size=9, color='808080')
F_HDR = Font(name=HEI, size=10, bold=True, color='FFFFFF')
F_SEC = Font(name=HEI, size=10, bold=True, color='1F4E79')
F_LBL = Font(name=HEI, size=9, bold=True, color='833C00')
F_IN = Font(name=SONG, size=10, bold=True, color='1F4E79')
F_AUTO = Font(name=SONG, size=10, color='006100')
F_AUTOB = Font(name=SONG, size=10, bold=True, color='006100')
F_TOT = Font(name=HEI, size=10, bold=True, color='000000')
F_TOTN = Font(name=SONG, size=10, bold=True, color='000000')
F_SUBL = Font(name=HEI, size=10, bold=True, color='833C00')
F_SUBN = Font(name=SONG, size=10, bold=True, color='C00000')
F_CHK = Font(name=HEI, size=10, bold=True, color='C00000')
F_HELP = Font(name=SONG, size=8, color='BFBFBF')

def fill(c): return PatternFill('solid', fgColor=c)
FL_TITLE = fill('1F4E79'); FL_HDR = fill('2E75B6'); FL_SEC = fill('9DC3E6')
FL_AUTO = fill('E2EFDA'); FL_IN = fill('DDEBF7'); FL_LBL = fill('FCE4D6')
FL_TOT = fill('9DC3E6'); FL_TOTM = fill('FCE4D6'); FL_SUB = fill('FFC000')
FL_NOTE = fill('FFF2CC'); FL_NONE = PatternFill(fill_type=None)

AC = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL = Alignment(horizontal='left', vertical='center', wrap_text=True)
AR = Alignment(horizontal='right', vertical='center')
ACN = Alignment(horizontal='center', vertical='center')
_t = Side(style='thin', color='BFBFBF')
BOX = Border(left=_t, right=_t, top=_t, bottom=_t)
NOB = Border()

MONEY = '#,##0.00;[Red]\\-#,##0.00;\\-'
QTY = '#,##0.###;[Red]\\-#,##0.###;\\-'
INT = '#,##0;[Red]\\-#,##0;\\-'
DATE = 'yyyy/m/d'
MONTH = 'yyyy"年"m"月"'
INF = 2958465          # 9999-12-31，「不限」的截止日
YEARS = '"2024,2025,2026,2027,2028,2029,2030"'   # 年份下拉（也可以直接手填别的年份）

# 资金日记账数据区
JR0, JR1 = 4, 603
def J(col):
    return f'资金日记账!${col}${JR0}:${col}${JR1}'

# 往来业务明细数据区（对账补充行扩到 300 行后，末行 12304）
WR0, WR1 = 4, 12304
def W(col):
    return f'往来业务明细!${col}${WR0}:${col}${WR1}'


def cell(ws, ref, value=None, font=F_AUTO, fl=FL_AUTO, al=AC, fmt=None, border=BOX):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font
    c.fill = fl
    c.alignment = al
    c.border = border
    if fmt:
        c.number_format = fmt
    return c


def wipe(ws):
    """清空一张表（保留表本身、位置、标签颜色以外的一切都清掉）"""
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    for row in ws.iter_rows():
        for c in row:
            c.value = None
            c.style = 'Normal'
    ws._cells.clear()
    ws.data_validations.dataValidation = []
    ws.conditional_formatting._cf_rules.clear()
    ws.column_dimensions.clear()
    ws.row_dimensions.clear()
    ws.freeze_panes = None
    ws.auto_filter.ref = None
    ws.print_area = None
    ws._print_rows = None
    ws._hyperlinks = []
    ws.sheet_view.showGridLines = False


def title(ws, text, c1, c2, row=1, height=30):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    cell(ws, f'{c1}{row}', text, F_TITLE, FL_TITLE, AC, border=NOB)
    ws.row_dimensions[row].height = height


def note(ws, text, c1, c2, row, height=30, warn=False):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    cell(ws, f'{c1}{row}', text, Font(name=HEI, size=9, bold=warn, color='833C00' if warn else '808080'),
         FL_NOTE if warn else FL_NONE, AL, border=NOB)
    ws.row_dimensions[row].height = height


def section(ws, text, c1, c2, row, height=20):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    cell(ws, f'{c1}{row}', text, F_SEC, FL_SEC, AL)
    for i in range(CI(c1) + 1, CI(c2) + 1):
        ws.cell(row=row, column=i).border = BOX
    ws.row_dimensions[row].height = height


def header(ws, row, col1, labels, height=30):
    for i, t in enumerate(labels):
        cell(ws, f'{L(CI(col1) + i)}{row}', t, F_HDR, FL_HDR, AC)
    ws.row_dimensions[row].height = height


def widths(ws, spec):
    for k, v in spec.items():
        ws.column_dimensions[k].width = v


def hide(ws, *cols):
    for c in cols:
        ws.column_dimensions[c].hidden = True


def period(ws, row=2, col='A', month=False, year=2026, mon=None, open_ended=False, note_text=None, clamp=None):
    """时段选择块。返回 dict：start / end（实际起止日期格，绝对引用）、Y（年份格）、M（月份格或 None）。
       规则：起始日/截止日填了就用；否则有月份＝该月；否则＝全年。
       open_ended=True 时年份也留空＝不限（对账单用）。"""
    c0 = CI(col)
    refs = {}
    labels = [('年　份', 'Y', year)]
    if month:
        labels.append(('月　份', 'M', mon))
    labels += [('起始日', 'S', None), ('截止日', 'E', None)]
    k = c0
    for lab, key, dflt in labels:
        cell(ws, f'{L(k)}{row}', lab, F_LBL, FL_LBL, AC)
        c = cell(ws, f'{L(k + 1)}{row}', dflt, F_IN, FL_IN, AC,
                 fmt=DATE if key in 'SE' else '0')
        refs[key] = f'${L(k + 1)}${row}'
        k += 2
    Y, S, E = refs['Y'], refs['S'], refs['E']
    M = refs.get('M')
    if open_ended:
        start = f'=IF(N({S})>0,{S},IF(N({Y})>0,DATE({Y},{"IF(N(" + M + ")>0," + M + ",1)" if M else "1"},1),1))'
        end = (f'=IF(N({E})>0,{E},IF(N({Y})>0,' +
               (f'IF(N({M})>0,DATE({Y},{M}+1,0),DATE({Y},12,31))' if M else f'DATE({Y},12,31)') + f',{INF}))')
    else:
        # 年份留空＝资金日记账里最后一笔的年份（不用 TODAY()，免得整表每次编辑都重算）
        yy = f'IF(N({Y})>0,{Y},YEAR(MAX(资金日记账!$B$4:$B$603)))'
        start = f'=IF(N({S})>0,{S},' + (f'IF(N({M})>0,DATE({yy},{M},1),DATE({yy},1,1)))' if M else f'DATE({yy},1,1))')
        end = f'=IF(N({E})>0,{E},' + (f'IF(N({M})>0,DATE({yy},{M}+1,0),DATE({yy},12,31)))' if M else f'DATE({yy},12,31))')
    if clamp:
        # 按月矩阵只排 clamp 个月：截止日最多到起始月往后 clamp 个月的月底；
        # 只填了起始日、没填截止日 → 从起始日往后排满 clamp 个月（如 8 月～次年 7 月一个果季）
        win = f'DATE(YEAR({L(k + 1)}{row}),MONTH({L(k + 1)}{row})+{clamp},0)'
        end = f'=IF(N({E})>0,MIN({E},{win}),IF(N({S})>0,{win},MIN({end[1:]},{win})))'
    cell(ws, f'{L(k)}{row}', '实际统计', F_LBL, FL_LBL, AC)
    cell(ws, f'{L(k + 1)}{row}', start, F_AUTOB, FL_AUTO, AC, fmt=f'[>={INF}]"不限";[<=1]"不限";{DATE}')
    cell(ws, f'{L(k + 2)}{row}', '至', F_LBL, FL_LBL, AC)
    cell(ws, f'{L(k + 3)}{row}', end, F_AUTOB, FL_AUTO, AC, fmt=f'[>={INF}]"不限";[<=1]"不限";{DATE}')
    refs['start'] = f'${L(k + 1)}${row}'
    refs['end'] = f'${L(k + 3)}${row}'
    refs['last_col'] = L(k + 3)
    ws.row_dimensions[row].height = 24
    # 数据验证
    dvy = DataValidation(type='list', formula1=YEARS, allow_blank=True, showErrorMessage=False)
    dvy.add(Y.replace('$', ''))
    ws.add_data_validation(dvy)
    if M:
        dvm = DataValidation(type='list', formula1='"1,2,3,4,5,6,7,8,9,10,11,12"', allow_blank=True)
        dvm.add(M.replace('$', ''))
        ws.add_data_validation(dvm)
    dvd = DataValidation(type='date', operator='between', formula1='36526', formula2='73050', allow_blank=True,
                         showErrorMessage=True, errorTitle='日期', error='请填日期，如 2026/9/1')
    dvd.add(S.replace('$', ''))
    dvd.add(E.replace('$', ''))
    ws.add_data_validation(dvd)
    return refs


def period_text(p):
    """“2026/1/1 至 2026/12/31” 文字（给说明行用）"""
    s, e = p['start'], p['end']
    return (f'IF({s}<=1,"不限",TEXT({s},"yyyy/m/d"))&" 至 "&IF({e}>={INF},"不限",TEXT({e},"yyyy/m/d"))')


def in_period(datecol, p):
    """SUMIFS 用的日期条件（起止都含）"""
    return f'{datecol},">="&{p["start"]},{datecol},"<="&{p["end"]}'


def pool(ws, col_key, col_first, col_cnt, key_expr, r0=JR0, r1=JR1, hdr_row=3, label='去重'):
    """在 ws 的三列隐藏辅助列里，对齐资金日记账第 r0..r1 行，做「按条件去重、按先后编号」：
       col_key 放候选值（不满足条件为 ""），col_first 标首次出现，col_cnt 累计编号。
       key_expr 是一个带 {r} 占位的表达式（资金日记账行号）。返回 (键列区域, 计数列区域, 总数格)。"""
    for c in (col_key, col_first, col_cnt):
        ws[f'{c}{hdr_row}'] = label
        ws[f'{c}{hdr_row}'].font = F_HELP
    for r in range(r0, r1 + 1):
        ws[f'{col_key}{r}'] = '=' + key_expr.format(r=r)
        ws[f'{col_first}{r}'] = (f'=IF({col_key}{r}="",0,IF(MATCH({col_key}{r},${col_key}${r0}:${col_key}${r1},0)'
                                 f'={r - r0 + 1},1,0))')
        ws[f'{col_cnt}{r}'] = f'=N({col_cnt}{r - 1})+{col_first}{r}' if r > r0 else f'={col_first}{r}'
        for c in (col_key, col_first, col_cnt):
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, col_key, col_first, col_cnt)
    return (f'${col_key}${r0}:${col_key}${r1}', f'${col_cnt}${r0}:${col_cnt}${r1}', f'${col_cnt}${r1}')


def list_item(keys, cnts, total, k_expr):
    """第 k 个去重值"""
    return f'IF({k_expr}>{total},"",INDEX({keys},MATCH({k_expr},{cnts},0)))'


def dv_list(ws, ref, formula, err=False):
    dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=err)
    dv.add(ref)
    ws.add_data_validation(dv)
