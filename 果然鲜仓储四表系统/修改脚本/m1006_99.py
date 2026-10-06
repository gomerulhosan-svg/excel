# -*- coding: utf-8 -*-
"""10/06 这一轮的收尾小修（最后一遍复核挑出来的显示问题），在各本自己的模块之后跑：

  《01》库存结余 H、M 列宽 13→15（合计行 1,492,785.00 / 1,199,361.17 两位小数加粗会显示 #####）；
       公司购买接口 A2 行高加够；库存总结余 A2 补一句「货主＝果然鲜的自购成品不在这张表」。
  《03》库存价值与欠款比对 E 列宽 13→15（合计 1,047,358.29）；对接源_01库存 A2 行高加够。
  《04》物料与筐子销售对账单、公司购买果品对账单 A3 说明开自动换行、行高加够（这一轮补的话看不到）。
"""
import copy
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L


def set_width(ws, col, width):
    """给单列设宽：先把包含这一列的列组拆开（拆成左段/本列/右段），避免 <col> 重叠"""
    i = CI(col)
    for key, d in list(ws.column_dimensions.items()):
        lo, hi = d.min or CI(key), d.max or CI(key)
        if lo <= i <= hi and lo != hi:
            del ws.column_dimensions[key]
            for a, b in ((lo, i - 1), (i, i), (i + 1, hi)):
                if a > b:
                    continue
                nd = copy.copy(d)
                nd.index, nd.min, nd.max = L(a), a, b
                ws.column_dimensions[L(a)] = nd
            break
    ws.column_dimensions[col].width = width


def wrap(ws, ref, height):
    c = ws[ref]
    al = copy.copy(c.alignment)
    c.alignment = Alignment(horizontal=al.horizontal or 'left', vertical=al.vertical or 'center', wrap_text=True,
                            indent=al.indent)
    r = c.row
    if (ws.row_dimensions[r].height or 0) < height:
        ws.row_dimensions[r].height = height


def apply_01(wb):
    ws = wb['库存结余']
    set_width(ws, 'H', 15)
    set_width(ws, 'M', 15)
    wrap(wb['公司购买接口'], 'A2', 48)
    ws = wb['库存总结余']
    v = ws['A2'].value
    add = '注意：货主＝果然鲜的公司自购成品不在这张表（只列有原料入库的货主），所以在库成品合计会比【库存结余】少一点。'
    if isinstance(v, str) and add not in v:
        ws['A2'].value = v + add
        wrap(ws, 'A2', 75)


def apply_03(wb):
    set_width(wb['库存价值与欠款比对'], 'E', 15)
    wrap(wb['对接源_01库存'], 'A2', 48)


def apply_04(wb):
    wrap(wb['物料与筐子销售对账单'], 'A3', 55)
    wrap(wb['公司购买果品对账单'], 'A3', 55)
