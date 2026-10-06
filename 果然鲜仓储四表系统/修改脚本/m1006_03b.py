# -*- coding: utf-8 -*-
"""10/06 这一轮（A050 1006）《03 财务账套与报表》报表部分（03B）。集成时在 m1006_03a 之后跑。

用法：wb = openpyxl.load_workbook(底稿)；m1006_03a.apply(wb)；apply(wb)；wb.save(成品)。
对底稿重复跑结果一样（版式、公式按固定规则重写；用户手填的值从底稿原位置读出来再放到新位置）。
单独对底稿跑（不跑 03A）也行：公式按 K9/K10/K11 的新范围写，那些行在底稿里还空着，只是取到 0。

这一轮改了什么（对应 A050 1006 实施规格第 4 节 03B）：

① 【对接源_01库存】按 K8 加三列：L 在库成品吨位（含损耗）KG ← [1]库存等级接口 J，M 成品出库总重KG ← K，
   N 损耗KG ← L；K 列（平均重量）公式不变。原来 L1:M1 的「有效行数」计数格挪到 R1:S1。

② 【库存价值与欠款比对】不插列，按段重写（上下两段共用一套列字母，插列会把两段一起挪）：
   · 第 3 行加蓝色参数格 D3「净毛比」＝0.87（没填单箱净重时用）。
   · 上半段「一、按客户比对」：表头第 5 行，数据 6～65 扩到 6～205（200 行，跟 _自动清单 的 200 个单位对齐），
     合计第 206 行。列：A 序号｜B 客户｜C 库存笔数｜D 结余数量｜E 结余净重｜F 结余成品吨位净重（＝下半段该客户
     K 合计）｜G 在库成品毛重（＝下半段 J 合计）｜H 库存价值｜I 未发成品价值（＝下半段 M 合计）｜J 应收款余额｜
     K 差额＝H＋I－J｜L 建议｜M 说明（不换行，文字往右边空格里溢出，不占下半段的列宽）。条件格式跟到 K 列。
   · 下半段「二、按客户＋等级明细」整体往下挪 141 行：区块标题第 208 行，表头第 209 行，数据 210～609
     （↔ 对接源第 4～403 行，按位置 INDEX 取），合计第 610 行。列：…｜G 结余净重＝结余数量×（P 填了就用 P，
     否则 Q×净毛比）｜H 库存价值＝G×I｜I 单价（手填）｜J 在库成品吨位毛重（← 对接源 L）｜K 估算在库净重＝J×净毛比
     （净毛比＝P÷Q，没填 P 用 D3）｜L 成品单价（手填，元/KG）｜M 未发成品价值＝K×L｜N 库存状态｜O 备注｜
     P 单箱净重（手填）｜Q 平均重量（← 对接源 K）。筛选 A209:Q610，合并区、列宽、冻结（C6）跟着改。
   · 用户手填的值原样搬家：I69:I468 的单价 → I210:I609（同列，行号＋141）；K69:K468 里手打的单箱净重
     （K76:K137 共 54 格）→ P210:P609（行号＋141）。手填 0 也照 0 算（ISNUMBER 判断「填没填」）。
   · 用户的 N16:P16 留在原位（不跟新列冲突）；O16 原来 =I16+I44+I47 加的是应收款，应收款挪到 J 列后跟着改成
     =J16+J44+J47（相当于 Excel 插列时的自动调整），N16、P16 原样不动；N16 加批注说明。

③ 【果然鲜总表】在原位置重建（保留她的四块栏目和标签）：
   · 第 2 行时段（年份／起始日／截止日）；标题带年份。
   · 块一「收入（按客户）」：本期有收入的单位自动列出（最多 100 家，超了有提示），最上面、最下面各一行合计
     （合计直接按源表算，不受显示行数限制）：仓储费／周转费／装卸费（往来业务明细「03计价」；计价日期全空，
     只选年份时连没填日期的一起算）、筐子销售、发泡网（品名在第 4 行蓝格，默认 香梨网、苹果网）、包装物料、
     优惠（对账补充行摘要含「优惠/抹零」）、小计；B、C、J 三列她原表表头空着，先留空待定。
   · 块二标题改「其他收入」并注明原叫「营业外收入」：筐子／托盘租赁费（02使用费）、扎伤梨款／扎伤苹果款／垃圾费
     （日记账收入行的项目）、盘盈款（没有来源，写「待定」）、果然鲜销售果品款（01果品 应收，按购买方）、利息、其他
     （其他服务费）。客户同样自动列出（最多 60 家）。
   · 块三改成按月 12 行＋「其余月份」1 行（区间超过 12 个月也不漏），保留她的成本栏目：所有费用成本（果然鲜类，
     不含右边单列的福利/商务/送礼）、装卸费成本（02装筐费 应付；装筐费单价没设之前按日记账「装卸费支出」）、
     福利、商务、送礼、筐子成本（卖出数量×该筐采购均价）、发泡网成本／包装物料成本（K11：对接源_02物料 W:AB，
     月末一次加权）、采购果品（01果品 应付，含原料公司购买）、赔付、小计。
   · 块四全公式：公司利润、销售果品利润、发泡网利润、筐子利润、包装物料利润、装卸利润、营业外利润，
     「公司收／个人收」分开合计（放在她原来写「帮我统计个人拿多少利润」「公司拿多少利润」的旁边），
     加「与利润表对账」一行（利润表期间跟本表不一样时提示）和本表收支自洽核对。
   · 表底「说明」列出取不到／有前提的格子和原因，以及要问她的 6 个问题。
   · 取数范围按接口约定：往来业务明细 $4:$17603（K9），资金日记账 $4:$4003（K10），对接源_02物料 W4:AB27（K11）。

④ 【主页】最后另起一块「本轮（10/6）新增/改动 · 报表」：果然鲜总表、库存价值与欠款比对、对接源_01库存 三个入口；
   改两条旧说明；【使用说明】末尾补两条（接在 03A 的「账务」几条后面）。
⑤ 【杨洋欠款】不动。

复核后修正（10/6 第二轮）：
· 果然鲜总表：重建前按格子文字记下她的批注（原 B46「公司利润」、H46「营业外利润」），重建后挂回同样文字的
  标签格（B201、H201），每格 new 一个 Comment；说明 6 补「单价为 0 的销售出库不算卖筐」；说明各行按字数给行高。
· 库存价值与欠款比对：清掉底稿遗留的筛选条件（filterColumn、sortState）；上半段她手工隐藏的客户行按原行号回放
  （客户顺序不变），下半段不回放；设冻结前先把 selection 清成一个（不再堆出 6 个）。
· 二表 O 备注加两支：原料已出完（F＝0）、J＞0 且《01》损耗≠0 →「原料已出完：这 x KG 在《01》按方案 A 算作损耗…」；
  没有原料入库（F＝0、Q＝0、J≠0，扎伤/果然鲜成品组合）→「这个等级没有原料入库，只有成品出库，负数正常…」；
  负数那支补「填成品单价时负数等级按同价一起填，正负才抵得掉」。O 列放宽到 60。
· 一表：应收款＜0（我方欠对方）的单位 K 差额留空；J 合计只加＞0 的，L/M 另列「我方欠对方合计」，N:Q 注明口径。
· N4 加跳到二表的链接，二表标题右边（P:Q）加回顶部的链接；A2 说明补「按货主＋品种统一填价」「下半段从第 208 行开始」。
"""
import copy
import re

from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.views import Selection

from common0928 import (F_TITLE, F_NOTE, F_HDR, F_SEC, F_LBL, F_IN, F_AUTO, F_AUTOB, F_TOT, F_TOTN, F_CHK, F_HELP,
                        FL_TITLE, FL_HDR, FL_SEC, FL_AUTO, FL_IN, FL_LBL, FL_TOT, FL_NOTE, FL_NONE,
                        AC, AL, AR, ACN, BOX, NOB, MONEY, QTY, INT, DATE, MONTH,
                        cell, wipe, title, note, section, header, widths, hide, period)

# ── 接口约定的范围（K9 / K10 / K11）──────────────────────────────────────────
WL_END = 17603                    # 往来业务明细 4～17603
RJ_END = 4003                     # 资金日记账 4～4003
WL = lambda c: f'往来业务明细!${c}$4:${c}${WL_END}'
RJ = lambda c: f'资金日记账!${c}$4:${c}${RJ_END}'
MAT = lambda c: f'对接源_02物料!${c}$4:${c}$27'          # K11：W 月份｜Y 发泡网成本｜AA 其他包装物料成本

# ── 库存价值与欠款比对 新版式 ────────────────────────────────────────────────
SHEET_LV = '库存价值与欠款比对'
SRC = '对接源_01库存'
INV_N = 400                       # 对接源 4～403
U_HDR, U0, U1, UT = 5, 6, 205, 206              # 上半段：表头、数据、合计
L_SEC, L_HDR, L0, L1, LT = 208, 209, 210, 609, 610   # 下半段
OLD_L0, OLD_L1 = 69, 468          # 底稿里下半段数据行
SHIFT = L0 - OLD_L0               # 141
PARAM = '$D$3'                    # 净毛比参数格

WEIGHT = '#,##0.00;[Red]\\-#,##0.00;\\-'
AL_NW = Alignment(horizontal='left', vertical='center', wrap_text=False)
F_LINK = Font(name='微软雅黑', size=9, bold=True, color='0563C1', underline='single')


def _freeze(ws, ref):
    """设冻结前先把 selection 清成一个，免得 openpyxl 在旧的 4 个 selection 上再插 2 个（规范上限 4 个、pane 不能重复）"""
    ws.sheet_view.selection = [Selection()]
    ws.freeze_panes = ref


def _is_formula(v):
    return isinstance(v, str) and v.startswith('=')


def _link_prefix(src):
    """对接源里引用《01》用的外链序号，如 [1]"""
    m = re.match(r"=IFERROR\(IF\((\[\d+\])库存等级接口!", src['H4'].value or '')
    if not m:
        raise SystemExit('对接源_01库存!H4 不是预期的外链公式，停下来查一下')
    return m.group(1)


# ───────────────────────────── ① 对接源_01库存 ─────────────────────────────
def source_01(wb):
    ws = wb[SRC]
    link = _link_prefix(ws)
    # 计数格 L1:M1 → R1:S1（核实过除主页超链接外没人引用）
    if ws['L1'].value == '有效行数':
        for a, b in (('L1', 'R1'), ('M1', 'S1')):
            ws[b].value = ws[a].value
            ws[b]._style = copy.copy(ws[a]._style)
            ws[a].value = None
            ws[a].style = 'Normal'
    ws.column_dimensions['R'].width = 9
    ws.column_dimensions['S'].width = 9
    cols = (('L', 'J', '在库成品吨位\n（含损耗）KG'), ('M', 'K', '成品出库总重KG'), ('N', 'L', '损耗KG'))
    for col, scol, head in cols:
        ws[f'{col}3'] = head
        ws[f'{col}3']._style = copy.copy(ws['K3']._style)
        ws[f'{col}3'].alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for i in range(INV_N):
            r = 4 + i
            ref = f'{link}库存等级接口!${scol}${r}'
            c = ws[f'{col}{r}']
            c.value = f'=IFERROR(IF({ref}=0,"",{ref}),"")'
            c._style = copy.copy(ws[f'K{r}']._style)
            c.number_format = '#,##0.00'
        ws.column_dimensions[col].width = 14
    ws.row_dimensions[3].height = 30
    a2 = ws['A2'].value or ''
    if 'L/M/N' not in a2:
        ws['A2'] = a2 + ' K＝平均重量；L/M/N＝在库成品（含损耗）、成品出库总重、损耗（《01》库存等级接口 I～L，不随日期筛选）。'


# ───────────────────────────── ② 库存价值与欠款比对 ─────────────────────────────
def _grab_user_values(ws):
    """把用户手填的值读出来：{行偏移: 值}。认底稿版式（下半段 69～468）和本模块跑过的版式（210～609）。"""
    if isinstance(ws['A67'].value, str) and ws['A67'].value.startswith('二、'):
        r0, r1, pcol, lcol = OLD_L0, OLD_L1, 'K', None
        u1 = 65                                   # 底稿上半段数据 6～65
    elif isinstance(ws[f'A{L_SEC}'].value, str) and ws[f'A{L_SEC}'].value.startswith('二、'):
        r0, r1, pcol, lcol = L0, L1, 'P', 'L'
        u1 = U1
    else:
        raise SystemExit('【库存价值与欠款比对】版式认不出来，停下来查一下')
    # 上半段她手工隐藏的客户行：客户按 _自动清单 E 列顺序逐行 INDEX（ROW()-5），新旧版式同一行是同一家，原行号回放。
    # 下半段的隐藏是旧筛选（F 列结余数量）留下的，不回放（行号已变、筛选条件也清掉）。
    hidden = [r for r in range(U0, u1 + 1) if ws.row_dimensions[r].hidden]
    price, net, fprice = {}, {}, {}
    for r in range(r0, r1 + 1):
        i = r - r0
        v = ws[f'I{r}'].value
        if v is not None and not _is_formula(v):
            price[i] = v
        v = ws[f'{pcol}{r}'].value
        if v is not None and not _is_formula(v):
            net[i] = v
        if lcol:
            v = ws[f'{lcol}{r}'].value
            if v is not None and not _is_formula(v):
                fprice[i] = v
    user = {}
    for ref in ('N16', 'O16', 'P16'):
        c = ws[ref]
        if c.value is not None:
            user[ref] = (c.value, copy.copy(c._style))
    if 'O16' in user and user['O16'][0] == '=I16+I44+I47':      # 应收款从 I 挪到 J
        user['O16'] = ('=J16+J44+J47', user['O16'][1])
    param = ws['D3'].value if r0 == L0 and isinstance(ws['D3'].value, (int, float)) else 0.87
    return price, net, fprice, user, param, hidden


def stock_compare(wb):
    ws = wb[SHEET_LV]
    price, net, fprice, user, param, hidden = _grab_user_values(ws)
    grid = ws.sheet_view.showGridLines
    wipe(ws)
    ws.sheet_view.showGridLines = grid

    # 标题、说明、参数
    title(ws, '库 存 价 值 与 欠 款 比 对', 'A', 'Q')
    note(ws, '★ 淡蓝格是手工填的：下半段 I「单价」（原料，元/KG）、L「成品单价」（元/KG）、P「单箱净重」（KG/件），'
             '第 3 行 D3「净毛比」。下半段每行：结余净重＝结余数量×单箱净重（P 没填就用 平均重量×净毛比），'
             '库存价值＝结余净重×单价；在库成品毛重取《01》库存等级接口（加工出库的原料重－成品已发出的重量，含损耗），'
             '估算在库净重＝毛重×净毛比（净毛比＝P÷Q，P 没填用 D3），未发成品价值＝估算净重×成品单价。'
             '未发成品价值请按「货主＋品种」统一填价（同一货主同一品种的各等级填同一个成品单价，重新分级造成的正负才抵得掉）。'
             '上半段按客户合计：差额＝库存价值＋未发成品价值－应收款（应收款取【往来对账单明细】期末余额）；'
             '应收款为负的单位（借款人、出借方、押金等我方欠对方的）差额留空，合计行的应收款也不含它们。'
             '注意：手填的值是跟着行走的，《01》里删行、改历史行的货主/等级会让顺序变，填完请核对一下客户和等级。'
             f'下半段从第 {L_SEC} 行开始。',
         'A', 'Q', 2, 72)
    # 跳到二表（A4:M4 是一表的区块标题，N4:Q4 空着）
    ws.merge_cells('N4:Q4')
    cell(ws, 'N4', f'=HYPERLINK("#\'{SHEET_LV}\'!A{L_SEC}","↓ 二表（客户＋等级明细，填单价/单箱净重）在第 {L_SEC} 行")',
         F_LINK, FL_NONE, AL_NW, border=NOB)
    ws.merge_cells('A3:C3')
    cell(ws, 'A3', '净毛比（P 没填单箱净重时用）', F_LBL, FL_LBL, AC)
    for c in 'BC':
        ws[f'{c}3'].border = BOX
    cell(ws, 'D3', param, F_IN, FL_IN, ACN, fmt='0.000')
    ws.merge_cells('E3:Q3')
    cell(ws, 'E3', f'=IF(AND(N({SRC}!$S$1)>0,COUNT({SRC}!$K$4:$K$403)=0),'
                   f'"⚠ 没接到《01》的平均重量/在库成品：四本放同一个文件夹，打开时选「更新链接」；《01》要用 10/6 以后的版本",'
                   f'"← 净毛比＝净重÷毛重。你手填的单箱净重 16/20/15 跟入库平均件重比，中位数约 0.874")',
         F_NOTE, FL_NONE, AL_NW, border=NOB)
    ws.row_dimensions[3].height = 22

    # ── 上半段 ──
    section(ws, '一、按 客 户 比 对（库存价值＋未发成品价值 ↔ 应收款 ↔ 差额）', 'A', 'M', 4)
    header(ws, U_HDR, 'A', ['序号', '客户', '库存笔数', '结余数量', '结余净重KG', '结余成品吨位\n净重KG（估算）',
                            '在库成品\n毛重KG', '库存价值\n（原料）', '未发成品价值', 'x', '差额（库存价值＋\n未发成品价值－应收）',
                            '建议', '说明'], 36)
    ws[f'J{U_HDR}'] = ('="应收款余额"&CHAR(10)&"（往来对账单明细·截至"&IF(往来对账单明细!$J$2>=2958465,"全部",'
                       'TEXT(往来对账单明细!$J$2,"yyyy/m/d"))&"）"')
    lb, le = f'$B${L0}', f'$B${L1}'
    sumif = lambda r, col: f'=IF($B{r}="","",ROUND(SUMIF({lb}:{le},$B{r},${col}${L0}:${col}${L1}),2))'
    for r in range(U0, U1 + 1):
        cell(ws, f'A{r}', f'=IF($B{r}="","",ROW()-{U0 - 1})', al=ACN)
        cell(ws, f'B{r}', f'=IFERROR(INDEX(_自动清单!$E$4:$E$203,ROW()-{U0 - 1}),"")', F_AUTOB, al=ACN)
        cell(ws, f'C{r}', f'=IF($B{r}="","",COUNTIF({SRC}!$B$4:$B$403,$B{r}))', al=AR, fmt=INT)
        cell(ws, f'D{r}', f'=IF($B{r}="","",SUMIF({SRC}!$B$4:$B$403,$B{r},{SRC}!$F$4:$F$403))', al=AR, fmt=INT)
        cell(ws, f'E{r}', sumif(r, 'G'), al=AR, fmt=WEIGHT)
        cell(ws, f'F{r}', sumif(r, 'K'), al=AR, fmt=WEIGHT)
        cell(ws, f'G{r}', sumif(r, 'J'), al=AR, fmt=WEIGHT)
        cell(ws, f'H{r}', sumif(r, 'H'), al=AR, fmt=MONEY)
        cell(ws, f'I{r}', sumif(r, 'M'), al=AR, fmt=MONEY)
        cell(ws, f'J{r}', f'=IF($B{r}="","",ROUND(IFERROR(INDEX(往来对账单明细!$F$7:$F$206,'
                          f'MATCH($B{r},往来对账单明细!$B$7:$B$206,0)),0),2))', al=AR, fmt=MONEY)
        # 应收款为负（我方欠对方：借款人、出借方、押金等）不算差额，留空；合计行自动不含
        cell(ws, f'K{r}', f'=IF($B{r}="","",IF(N($J{r})<0,"",ROUND(N($H{r})+N($I{r})-N($J{r}),2)))', al=AR, fmt=MONEY)
        cell(ws, f'L{r}', f'=IF($B{r}="","",IF(N($J{r})<0,"我方欠对方 · 不用催收",IF(N($J{r})=0,"没有欠款",'
                          f'IF(N($K{r})>=0,"货够抵账 · 可以先不收","货不够抵 · 该收钱了"))))', al=ACN)
        cell(ws, f'M{r}', f'=IF($B{r}="","",$B{r}&" 库里压着原料 "&TEXT(N($H{r}),"#,##0.00")&" 元、未发成品 "'
                          f'&TEXT(N($I{r}),"#,##0.00")&" 元，"&IF(N($J{r})<0,"我方还欠对方 "&TEXT(-N($J{r}),"#,##0.00")&" 元",'
                          f'"欠我们 "&TEXT(N($J{r}),"#,##0.00")&" 元，差额 "&TEXT(N($K{r}),"#,##0.00")&" 元"))',
             al=AL_NW)
        ws.row_dimensions[r].height = 18
    cell(ws, f'A{UT}', '合  计', F_TOT, FL_TOT, ACN)
    cell(ws, f'B{UT}', None, F_TOT, FL_TOT, ACN)
    for col, fmt in zip('CDEFGHIJK', (INT, INT, WEIGHT, WEIGHT, WEIGHT, MONEY, MONEY, MONEY, MONEY)):
        cell(ws, f'{col}{UT}', f'=ROUND(SUM({col}{U0}:{col}{U1}),2)', F_TOTN, FL_TOT, AR, fmt)
    # 应收款合计只加 >0 的（不含我方欠对方的单位）；我方欠对方的另列
    cell(ws, f'J{UT}', f'=ROUND(SUMIF(J{U0}:J{U1},">0"),2)', F_TOTN, FL_TOT, AR, MONEY)
    cell(ws, f'L{UT}', '我方欠对方合计', F_TOT, FL_TOT, ACN)
    cell(ws, f'M{UT}', f'=ROUND(-SUMIF(J{U0}:J{U1},"<0"),2)', F_TOTN, FL_TOT, AR, MONEY)
    ws.merge_cells(f'N{UT}:Q{UT}')
    cell(ws, f'N{UT}', '← J 合计只加应收款＞0 的，不含我方欠对方的单位（借款人、出借方、押金等，K 差额留空）；'
                       '我方欠对方的合计在左边 M 格', F_NOTE, FL_NONE, AL, border=NOB)
    ws.row_dimensions[UT].height = 30
    rng = f'K{U0}:K{U1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($B{U0}<>"",$K{U0}<0)'],
                                                   font=Font(b=True, color='FF9C0006'),
                                                   fill=PatternFill('solid', bgColor='FFFCE4E4', fgColor='FFFCE4E4')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND($B{U0}<>"",$K{U0}>=0,N($J{U0})>0)'],
                                                   fill=PatternFill('solid', bgColor='FFE2EFDA', fgColor='FFE2EFDA')))

    # ── 下半段 ──
    section(ws, '二、按 客 户 ＋ 等 级 明 细（淡蓝格手工填：I 单价、L 成品单价、P 单箱净重；其余自动）', 'A', 'O', L_SEC)
    ws.merge_cells(f'P{L_SEC}:Q{L_SEC}')
    cell(ws, f'P{L_SEC}', f'=HYPERLINK("#\'{SHEET_LV}\'!A1","↑ 回到顶部（一表）")', F_LINK, FL_SEC, ACN)
    ws[f'Q{L_SEC}'].border = BOX
    header(ws, L_HDR, 'A', ['序号', '客户', '品种', '等级', '单位', '结余数量', '结余净重KG\n（数量×单箱净重）', '库存价值',
                            '单价（手工填）\n元/KG', '在库成品吨位\n毛重KG（自动）', '估算在库\n净重KG', '成品单价\n（手工填）元/KG',
                            '未发成品价值', '库存状态', '备注', '单箱净重KG\n（手工填）', '平均重量\nKG/件（自动）'], 40)
    k = L0 - 1                      # ROW()-k ＝ 对接源 4～403 里的第几个（1～400）：第 210 行 ↔ 对接源第 4 行
    src = lambda col: f'INDEX({SRC}!${col}$4:${col}$403,ROW()-{k})'
    for r in range(L0, L1 + 1):
        i = r - L0
        b = f'$B{r}'
        cell(ws, f'A{r}', f'=IF({b}="","",ROW()-{L_HDR})', al=ACN)
        cell(ws, f'B{r}', f'=IF({src("B")}="","",{src("B")})', F_AUTOB, al=ACN)
        cell(ws, f'C{r}', f'=IF({b}="","",{src("C")})', al=ACN)
        cell(ws, f'D{r}', f'=IF({b}="","",{src("D")})', al=ACN)
        cell(ws, f'E{r}', f'=IF({b}="","",{src("E")})', al=ACN)
        cell(ws, f'F{r}', f'=IF({b}="","",N({src("F")}))', al=AR, fmt=INT)
        cell(ws, f'G{r}', f'=IF({b}="","",ROUND(N($F{r})*IF(ISNUMBER($P{r}),$P{r},N($Q{r})*{PARAM}),2))', al=AR, fmt=WEIGHT)
        cell(ws, f'H{r}', f'=IF({b}="","",ROUND(N($G{r})*N($I{r}),2))', al=AR, fmt=MONEY)
        cell(ws, f'I{r}', price.get(i), F_IN, FL_IN, ACN, fmt='0.00')
        cell(ws, f'J{r}', f'=IF({b}="","",N({src("L")}))', al=AR, fmt=WEIGHT)
        cell(ws, f'K{r}', f'=IF({b}="","",ROUND(N($J{r})*IF(AND(ISNUMBER($P{r}),N($Q{r})>0),$P{r}/$Q{r},{PARAM}),2))',
             al=AR, fmt=WEIGHT)
        cell(ws, f'L{r}', fprice.get(i), F_IN, FL_IN, ACN, fmt='0.00')
        cell(ws, f'M{r}', f'=IF({b}="","",ROUND(N($K{r})*N($L{r}),2))', al=AR, fmt=MONEY)
        cell(ws, f'N{r}', f'=IF({b}="","",{src("H")})', al=ACN)
        loss, outw = f'N({src("N")})', f'N({src("M")})'      # 《01》损耗KG、成品出库总重KG（对接源 N、M）
        cell(ws, f'O{r}', f'=IF({b}="","",IF(AND(N($F{r})<>0,N($I{r})=0),"← 还有原料结余，没填单价",'
                          f'IF(AND(N($F{r})<>0,NOT(ISNUMBER($P{r})),N($Q{r})=0),"← 没有入库件重，净重算不出",'
                          f'IF(AND(N($F{r})<>0,ISNUMBER($P{r}),N($P{r})=0),"← 单箱净重填的是 0，净重按 0 算",'
                          f'IF(AND(N($F{r})=0,N($Q{r})=0,N($J{r})<>0),"← 这个等级没有原料入库，只有成品出库，负数正常，按客户合计看",'
                          f'IF(N($J{r})<0,"← 在库成品为负：成品比加工出库的原料多（加工时可能重新分级），看客户合计；'
                          f'填成品单价时负数等级按同价一起填，正负才抵得掉",'
                          f'IF(AND(N($F{r})=0,N($J{r})>0,{loss}<>0),"← 原料已出完：这 "&TEXT($J{r},"#,##0.0")&'
                          f'" KG 在《01》按方案 A 算作损耗（损耗率 "&TEXT({loss}/(N($J{r})+{outw}),"0%")&'
                          f'"），多半是加工时重新分级；确认库里真有成品再填单价",'
                          f'IF(AND(N($J{r})>0,N($L{r})=0),"← 有在库成品，没填成品单价",""))))))))', al=AL_NW)
        cell(ws, f'P{r}', net.get(i), F_IN, FL_IN, ACN, fmt='0.###')
        cell(ws, f'Q{r}', f'=IF({b}="","",N({src("K")}))', al=AR, fmt='0.000')
        ws.row_dimensions[r].height = 18
    cell(ws, f'A{LT}', '合  计', F_TOT, FL_TOT, ACN)
    for col in 'BCDEFGHIJKLMNOPQ':
        cell(ws, f'{col}{LT}', None, F_TOT, FL_TOT, ACN)
    for col, fmt in zip('FGHJKM', (INT, WEIGHT, MONEY, WEIGHT, WEIGHT, MONEY)):
        cell(ws, f'{col}{LT}', f'=ROUND(SUM({col}{L0}:{col}{L1}),2)', F_TOTN, FL_TOT, AR, fmt)
    ws.row_dimensions[LT].height = 18
    ws.auto_filter.ref = f'A{L_HDR}:Q{LT}'
    ws.auto_filter.filterColumn = []          # 底稿的筛选条件（F 列结余数量的值清单）是旧版式的，清掉
    ws.auto_filter.sortState = None
    # 上半段她手工隐藏的客户行按原行号回放（客户顺序不变）
    for r in hidden:
        ws.row_dimensions[r].hidden = True

    # 用户的 N16:P16 留在原位
    for ref, (v, st) in user.items():
        ws[ref].value = v
        ws[ref]._style = st
    if 'N16' in user:
        ws['N16'].comment = Comment(
            '10/6 改表说明：上半段在「库存价值」后面加了一列「未发成品价值」(I 列)，应收款从 I 挪到 J，'
            '所以 O16 跟着从 =I16+I44+I47 改成 =J16+J44+J47；N16、P16 没动。\n'
            '注意：第 16/44/47 行现在是 张桂荣、伍元虎、李现东（客户顺序会随新单位变）。'
            '如果你想看的是热合曼一家合计，告诉我们改成按名字取。', '果然鲜系统')
        ws['N16'].comment.width = 320
        ws['N16'].comment.height = 140

    widths(ws, {'A': 6, 'B': 15, 'C': 11, 'D': 13, 'E': 13, 'F': 14, 'G': 14, 'H': 15, 'I': 15, 'J': 17,
                'K': 17, 'L': 20, 'M': 16, 'N': 12, 'O': 60, 'P': 13, 'Q': 13})
    _freeze(ws, 'C6')
    return {'price': len(price), 'net': len(net), 'fprice': len(fprice), 'user': sorted(user), 'hidden': len(hidden)}


# ───────────────────────────── ③ 果然鲜总表 ─────────────────────────────
ZB = '果然鲜总表'
N1, N2 = 100, 60                  # 块一、块二最多显示几家
# 行号
R_SEC1, R_G1, R_H1, R_T1 = 7, 8, 9, 10
R_D1 = R_T1 + 1                   # 11
R_E1 = R_D1 + N1 - 1              # 110
R_B1, R_W1 = R_E1 + 1, R_E1 + 2   # 111 底部合计、112 提示
R_SEC2 = R_W1 + 2                 # 114
R_G2, R_H2, R_T2 = R_SEC2 + 1, R_SEC2 + 2, R_SEC2 + 3   # 115/116/117
R_D2 = R_T2 + 1                   # 118
R_E2 = R_D2 + N2 - 1              # 177
R_B2, R_W2 = R_E2 + 1, R_E2 + 2   # 178/179
R_SEC3 = R_W2 + 2                 # 181
R_G3, R_H3, R_T3 = R_SEC3 + 1, R_SEC3 + 2, R_SEC3 + 3   # 182/183/184
R_M0 = R_T3 + 1                   # 185：12 个月
R_REST = R_M0 + 12                # 197：其余月份
R_B3 = R_REST + 1                 # 198
R_SEC4 = R_B3 + 2                 # 200
R_L4, R_TAG4, R_AMT4, R_HOW4 = R_SEC4 + 1, R_SEC4 + 2, R_SEC4 + 3, R_SEC4 + 4   # 201～204
R_CHK4, R_SELF4 = R_HOW4 + 2, R_HOW4 + 3        # 206/207
R_NOTE = R_SELF4 + 2              # 209
# 隐藏辅助
CAND0, CAND1, CANDW = 11, 410, 210  # V 列候选单位：11～210 取自往来业务明细，211～410 取自日记账收入单位
KZ0, KZ1 = 4, 3003                # AB:AD 筐子成本逐行（往来业务明细 02筐子 两段：2004～4004、15705～16703）


def _period_crit(col_range, s='$H$2', e='$J$2'):
    return f'{col_range},">="&{s},{col_range},"<="&{e}'


def _find_profit_cells(wb):
    """利润表：净利润（本期）那格、本期起止两格。找不到就用底稿位置。"""
    ws = wb['利润表']
    net, s, e = 'C17', 'H1', 'H2'
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 80), max_col=8):
        for c in row:
            v = c.value
            if not isinstance(v, str):
                continue
            if c.column == 1 and '净利润' in v and v.strip().startswith('四'):
                net = f'C{c.row}'
            elif v.strip() == '本期起':
                s = f'{L(c.column + 1)}{c.row}'
            elif v.strip() == '本期止':
                e = f'{L(c.column + 1)}{c.row}'
    return net, s, e


def _grab_comments(ws):
    """她在格子上写的批注：按格子文字记下来（重建后挂回同样文字的标签格）。{文字: (批注内容, 作者)}"""
    keep = {}
    for row in ws.iter_rows():
        for c in row:
            if c.comment is not None and isinstance(c.value, str) and c.value.strip():
                keep.setdefault(c.value.strip(), (c.comment.text, c.comment.author))
    return keep


def _put_comments(ws, keep, rows):
    """先在给定的行里找同样文字的格子，找不到再全表找；每格 new 一个 Comment（不能共用）。返回挂上的格子。"""
    done = []
    for text, (body, author) in keep.items():
        hit = None
        for r in list(rows) + list(range(1, ws.max_row + 1)):
            for c in range(1, 13):
                v = ws.cell(r, c).value
                if isinstance(v, str) and v.strip() == text:
                    hit = ws.cell(r, c)
                    break
            if hit is not None:
                break
        if hit is None:
            continue
        cm = Comment(body, author or '')
        cm.width, cm.height = 240, 80
        hit.comment = cm
        done.append(hit.coordinate)
    return done


def summary(wb):
    ws = wb[ZB]
    keep = _grab_comments(ws)
    wipe(ws)
    P = _period_crit(WL('C'))
    PR = _period_crit(RJ('B'))
    UND = '$T$2'                 # 1＝没填起止日，没填计价日期的仓储类收入也算进来

    # 标题、时段、说明、参数
    ws.merge_cells('A1:L1')
    cell(ws, 'A1', '=IF(N($B$2)>0,$B$2&"年度","")&"果然鲜仓储有限公司账目总明细"', F_TITLE, FL_TITLE, AC, border=NOB)
    ws.row_dimensions[1].height = 30
    p = period(ws, 2, 'A', year=2026, open_ended=True)
    assert p['start'] == '$H$2' and p['end'] == '$J$2', p
    note(ws, '只填第 2 行：年份＝看全年，起始日/截止日可以缩小到某段时间。块一、块二按【往来业务明细】（应收发生额，'
             '不是收到的钱）和【资金日记账】收入取；块三按【资金日记账】【往来业务明细】和《02》物料成本取；块四全部按'
             '前三块算。块一、块二的单位自动列出（本期有收入的才列），每块最上面、最下面各一行合计（合计按全部单位算）。',
         'A', 'L', 3, 36)
    cell(ws, 'A4', '发泡网品名', F_LBL, FL_LBL, AC)
    for c, v in (('B', '香梨网'), ('C', '苹果网'), ('D', None)):
        cell(ws, f'{c}4', v, F_IN, FL_IN, ACN)
    ws.merge_cells('E4:L4')
    cell(ws, 'E4', '← 块一「发泡网」＝这几个品名卖出去的；其余包装物料算「包装物料」。以后有新的发泡网品名填在蓝格里'
                   '（跟《02》【发泡网销售汇总】一个口径）。', F_NOTE, FL_NONE, AL, border=NOB)
    ws.row_dimensions[4].height = 30               # E4:L4 说明约 70 字，要两行
    ws.merge_cells('A5:L5')
    cell(ws, 'A5', f'=IF($T$5=0,"",IF({UND}=1,"⚠ 【原料入库计价】有 "&TEXT($T$5,"#,##0.00")&" 元仓储/周转/装卸费'
                   f'没填计价日期：现在没填起止日，这部分算进块一了；填了起止日（按月看）就算不进去。",'
                   f'"⚠ 【原料入库计价】有 "&TEXT($T$5,"#,##0.00")&" 元仓储/周转/装卸费没填计价日期，填了起止日时这部分'
                   f'没算进块一（仓储/周转/装卸费偏少）。"))', F_CHK, FL_NONE, AL, border=NOB)
    ws.row_dimensions[5].height = 28

    # 隐藏辅助：标量
    for r, lab, f in ((1, '辅助（勿改）', None),
                      (2, None, '=IF(AND(N($D$2)=0,N($F$2)=0),1,0)'),
                      (3, None, f'=MAX({WL("Q")})'),
                      (4, None, f'=IF(ROUND(SUMIFS({WL("K")},{WL("A")},"02装筐费"),2)<>0,1,0)'),
                      (5, None, f'=ROUND(SUMIFS({WL("K")},{WL("A")},"03计价",{WL("E")},"应收",{WL("C")},0),2)'),
                      (7, None, f'=IF($H$2>1,DATE(YEAR($H$2),MONTH($H$2),1),'
                                f'DATE(YEAR(MAX(1,MIN({RJ("B")}))),MONTH(MAX(1,MIN({RJ("B")}))),1))')):
        ws[f'T{r}'] = lab if lab else f
        ws[f'T{r}'].font = F_HELP
        if r == 7:
            ws['T7'].number_format = DATE
    ws['S2'], ws['S3'], ws['S4'], ws['S5'], ws['S7'] = '含未填日期', '往来单位数', '装筐费启用', '未填日期金额', '首月'
    for r in (2, 3, 4, 5, 7):
        ws[f'S{r}'].font = F_HELP

    # 隐藏辅助：候选单位（V）、块一标记/累计（W/X）、块二标记/累计（Y/Z）
    for c, lab in (('V', '候选单位'), ('W', '块一有数'), ('X', '块一累计'), ('Y', '块二有数'), ('Z', '块二累计')):
        ws[f'{c}{CAND0 - 1}'] = lab
        ws[f'{c}{CAND0 - 1}'].font = F_HELP
    aj = '_自动清单!$AJ$4:$AJ$203'
    for r in range(CAND0, CAND1 + 1):
        v = f'$V{r}'
        if r <= CANDW:
            ws[f'V{r}'] = f'=IF(ROW()-{CAND0 - 1}>$T$3,"",INDEX({WL("D")},MATCH(ROW()-{CAND0 - 1},{WL("Q")},0)))'
            base = f'{WL("K")},{WL("D")},{v}'
            ws[f'W{r}'] = (f'=IF({v}="",0,IF(ABS(SUMIFS({base},{WL("E")},"应收",{WL("A")},"03计价",{P})'
                           f'+{UND}*SUMIFS({base},{WL("E")},"应收",{WL("A")},"03计价",{WL("C")},0))'
                           f'+ABS(SUMIFS({base},{WL("E")},"应收",{WL("A")},"02筐子",{P}))'
                           f'+ABS(SUMIFS({base},{WL("E")},"应收",{WL("A")},"02包装",{P}))'
                           f'+ABS(SUMIFS({base},{WL("E")},"应收",{WL("A")},"03补充",{WL("F")},"*优惠*",{P}))'
                           f'+ABS(SUMIFS({base},{WL("E")},"应收",{WL("A")},"03补充",{WL("F")},"*抹零*",{P}))>0.004,1,0))')
            ws[f'X{r}'] = f'=N(X{r - 1})+W{r}' if r > CAND0 else f'=W{r}'
        else:
            j = f'ROW()-{CANDW}'
            ws[f'V{r}'] = (f'=IFERROR(IF(INDEX({aj},{j})="","",IF(COUNTIF($V${CAND0}:$V${CANDW},INDEX({aj},{j}))>0,"",'
                           f'INDEX({aj},{j}))),"")')
        rj = lambda col, val: 'ABS(SUMIFS(' + RJ('H') + ',' + RJ('E') + ',' + v + ',' + RJ(col) + ',"' + val + '",' + PR + '))'
        rjs = '+'.join(rj(col, val) for col, val in (('F', '利息'), ('F', '其他服务费'), ('G', '次果梨款'), ('G', '扎伤梨款'),
                                                     ('G', '次果苹果款'), ('G', '扎伤苹果款'), ('G', '垃圾费')))
        ws[f'Y{r}'] = (f'=IF({v}="",0,IF(ABS(SUMIFS({WL("K")},{WL("D")},{v},{WL("A")},"02使用费",{P}))'
                       f'+ABS(SUMIFS({WL("K")},{WL("D")},{v},{WL("A")},"01果品",{WL("E")},"应收",{P}))'
                       f'+{rjs}>0.004,1,0))')
        ws[f'Z{r}'] = f'=N(Z{r - 1})+Y{r}' if r > CAND0 else f'=Y{r}'
        for c in 'VWXYZ':
            ws[f'{c}{r}'].font = F_HELP
    TOT1, TOT2 = f'$X${CANDW}', f'$Z${CAND1}'

    # 隐藏辅助：筐子成本逐行（AB 往来业务明细行号、AC 卖出日期、AD 卖出数量×该筐采购均价）
    for c, lab in (('AB', '往来行号'), ('AC', '卖筐日期'), ('AD', '筐子成本')):
        ws[f'{c}{KZ0 - 1}'] = lab
        ws[f'{c}{KZ0 - 1}'].font = F_HELP
    rows = list(range(2004, 4005)) + list(range(15705, 16704))
    assert len(rows) == KZ1 - KZ0 + 1
    ix = lambda col, r: f'INDEX({WL(col)},$AB{r}-3)'
    for r, wr in zip(range(KZ0, KZ1 + 1), rows):
        ws[f'AB{r}'] = wr
        ws[f'AC{r}'] = (f'=IF(AND({ix("A", r)}="02筐子",{ix("E", r)}="应收",{ix("D", r)}<>""),{ix("C", r)},"")')
        same = f'{WL("A")},"02筐子",{WL("E")},"应付",{WL("G")},{ix("G", r)}'
        ws[f'AD{r}'] = (f'=IF($AC{r}="","",ROUND(N({ix("I", r)})*IFERROR(SUMIFS({WL("K")},{same})'
                        f'/SUMIFS({WL("I")},{same}),0),2))')
        for c in ('AB', 'AC', 'AD'):
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, 'S', 'T', 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD')

    # ── 块一 ──
    section(ws, '一、收 入（按客户 · 本期有收入的单位自动列出 · 往来口径＝应收发生额）', 'A', 'L', R_SEC1)
    grp = lambda r1, r2, c1, c2, text, fill=FL_HDR, font=F_HDR: _grp(ws, r1, r2, c1, c2, text, fill, font)
    grp(R_G1, R_H1, 'A', 'A', '客户/总收入')
    for c in 'BCJ':
        grp(R_G1, R_H1, c, c, '（空栏·待定）', fill=PatternFill('solid', fgColor='D9D9D9'),
            font=Font(name='微软雅黑', size=9, bold=True, color='595959'))
    grp(R_G1, R_G1, 'D', 'F', '主营业务收入')
    grp(R_G1, R_G1, 'G', 'I', '其他业务收入')
    grp(R_G1, R_H1, 'K', 'K', '优惠（负数）')
    grp(R_G1, R_H1, 'L', 'L', '小计')
    header(ws, R_H1, 'D', ['仓储费', '周转费', '装卸费', '筐子销售', '发泡网', '包装物料'], 24)
    ws.row_dimensions[R_G1].height = 22
    foam = lambda d: '+'.join(f'IF(${c}$4="",0,SUMIFS({WL("K")},{WL("A")},"02包装",{WL("E")},"应收",{d}{WL("G")},${c}$4,{P}))'
                              for c in 'BCD')

    def blk1(d):
        """d：按客户的附加条件（如 往来业务明细!$D,$A11,），合计行传空串"""
        jj = lambda f: (f'SUMIFS({WL("K")},{WL("A")},"03计价",{WL("E")},"应收",{d}{WL("F")},"{f}",{P})'
                        f'+{UND}*SUMIFS({WL("K")},{WL("A")},"03计价",{WL("E")},"应收",{d}{WL("F")},"{f}",{WL("C")},0)')
        return {
            'D': jj('仓储加工'), 'E': jj('周转打冷'), 'F': jj('装卸费'),
            'G': f'SUMIFS({WL("K")},{WL("A")},"02筐子",{WL("E")},"应收",{d}{P})',
            'H': foam(d),
            'I': f'SUMIFS({WL("K")},{WL("A")},"02包装",{WL("E")},"应收",{d}{P})-({foam(d)})',
            'K': (f'SUMIFS({WL("K")},{WL("A")},"03补充",{WL("E")},"应收",{d}{WL("F")},"*优惠*",{P})'
                  f'+SUMIFS({WL("K")},{WL("A")},"03补充",{WL("E")},"应收",{d}{WL("F")},"*抹零*",{P})'),
        }
    tot = blk1('')
    for rr in (R_T1, R_B1):
        cell(ws, f'A{rr}', '合　计', F_TOT, FL_TOT, ACN)
        for c in 'BCJ':
            cell(ws, f'{c}{rr}', None, F_TOTN, FL_TOT, AR)
        for c, f in tot.items():
            cell(ws, f'{c}{rr}', f'=ROUND({f},2)' if rr == R_T1 else f'={c}{R_T1}', F_TOTN, FL_TOT, AR, MONEY)
        cell(ws, f'L{rr}', f'=ROUND(SUM(D{rr}:K{rr}),2)', F_TOTN, FL_TOT, AR, MONEY)
    for r in range(R_D1, R_E1 + 1):
        a = f'$A{r}'
        cell(ws, f'A{r}', f'=IF(ROW()-{R_D1 - 1}>{TOT1},"",INDEX($V${CAND0}:$V${CANDW},'
                          f'MATCH(ROW()-{R_D1 - 1},$X${CAND0}:$X${CANDW},0)))', F_AUTOB, al=ACN)
        for c in 'BCJ':
            cell(ws, f'{c}{r}', None, al=AR, fmt=MONEY)
        for c, f in blk1(f'{WL("D")},{a},').items():
            cell(ws, f'{c}{r}', f'=IF({a}="","",ROUND({f},2))', al=AR, fmt=MONEY)
        cell(ws, f'L{r}', f'=IF({a}="","",ROUND(SUM(D{r}:K{r}),2))', F_AUTOB, al=AR, fmt=MONEY)
    ws.merge_cells(f'A{R_W1}:L{R_W1}')
    cell(ws, f'A{R_W1}', f'=IF({TOT1}>{N1},"⚠ 本期有收入的单位共 "&{TOT1}&" 家，只列出前 {N1} 家；合计行按全部单位算",'
                         f'IF(ROUND(SUM(L{R_D1}:L{R_E1})-L{R_T1},2)=0,"本期有收入的单位共 "&{TOT1}&" 家，全部列出 ✔ 明细合计＝合计",'
                         f'"⚠ 明细合计比合计差 "&TEXT(L{R_T1}-SUM(L{R_D1}:L{R_E1}),"#,##0.00")&"（有收入的行没填往来单位）"))',
         F_CHK, FL_NONE, AL, border=NOB)

    # ── 块二 ──
    section(ws, '二、其 他 收 入（你原表叫「营业外收入」：筐子/托盘租赁、扎伤果款、自营果品款按会计都算经营收入，'
                '所以改叫其他收入；利息会计上是冲财务费用）', 'A', 'L', R_SEC2)
    grp(R_G2, R_H2, 'A', 'A', '客户/总收入')
    grp(R_G2, R_G2, 'B', 'J', '其他收入（原叫营业外收入）')
    grp(R_G2, R_H2, 'K', 'K', '小计')
    header(ws, R_H2, 'B', ['筐子租赁费', '托盘租赁费', '扎伤梨款', '扎伤苹果款', '垃圾费', '盘盈款', '果然鲜销售\n果品款',
                           '利息', '其他'], 30)
    ws.row_dimensions[R_G2].height = 22
    rjh = lambda e, crit: f'SUMIFS({RJ("H")},{e}{crit},{PR})'

    def blk2(d, e):
        return {
            'B': f'SUMIFS({WL("K")},{WL("A")},"02使用费",{d}{WL("G")},"<>*托盘*",{P})',
            'C': f'SUMIFS({WL("K")},{WL("A")},"02使用费",{d}{WL("G")},"*托盘*",{P})',
            'D': rjh(e, f'{RJ("G")},"次果梨款"') + '+' + rjh(e, f'{RJ("G")},"扎伤梨款"'),
            'E': rjh(e, f'{RJ("G")},"次果苹果款"') + '+' + rjh(e, f'{RJ("G")},"扎伤苹果款"'),
            'F': rjh(e, f'{RJ("G")},"垃圾费"'),
            'H': f'SUMIFS({WL("K")},{WL("A")},"01果品",{WL("E")},"应收",{d}{P})',
            'I': rjh(e, f'{RJ("F")},"利息"'),
            'J': rjh(e, f'{RJ("F")},"其他服务费"'),
        }
    tot = blk2('', '')
    for rr in (R_T2, R_B2):
        cell(ws, f'A{rr}', '合　计', F_TOT, FL_TOT, ACN)
        for c, f in tot.items():
            cell(ws, f'{c}{rr}', f'=ROUND({f},2)' if rr == R_T2 else f'={c}{R_T2}', F_TOTN, FL_TOT, AR, MONEY)
        cell(ws, f'G{rr}', '待定', F_TOTN, FL_TOT, ACN)
        cell(ws, f'K{rr}', f'=ROUND(SUM(B{rr}:J{rr}),2)', F_TOTN, FL_TOT, AR, MONEY)
        cell(ws, f'L{rr}', None, F_TOTN, FL_NONE, AR, border=NOB)
    for r in range(R_D2, R_E2 + 1):
        a = f'$A{r}'
        cell(ws, f'A{r}', f'=IF(ROW()-{R_D2 - 1}>{TOT2},"",INDEX($V${CAND0}:$V${CAND1},'
                          f'MATCH(ROW()-{R_D2 - 1},$Z${CAND0}:$Z${CAND1},0)))', F_AUTOB, al=ACN)
        for c, f in blk2(f'{WL("D")},{a},', f'{RJ("E")},{a},').items():
            cell(ws, f'{c}{r}', f'=IF({a}="","",ROUND({f},2))', al=AR, fmt=MONEY)
        cell(ws, f'G{r}', f'=IF({a}="","","待定")', al=ACN)
        cell(ws, f'K{r}', f'=IF({a}="","",ROUND(SUM(B{r}:J{r}),2))', F_AUTOB, al=AR, fmt=MONEY)
    ws.merge_cells(f'A{R_W2}:K{R_W2}')
    cell(ws, f'A{R_W2}', f'=IF({TOT2}>{N2},"⚠ 本期有其他收入的单位共 "&{TOT2}&" 家，只列出前 {N2} 家；合计行按全部单位算",'
                         f'IF(ROUND(SUM(K{R_D2}:K{R_E2})-K{R_T2},2)=0,"本期有其他收入的单位共 "&{TOT2}&" 家，全部列出 ✔ 明细合计＝合计",'
                         f'"⚠ 明细合计比合计差 "&TEXT(K{R_T2}-SUM(K{R_D2}:K{R_E2}),"#,##0.00")&"（日记账收入行没填往来单位，或单位超过 200 家）"))',
         F_CHK, FL_NONE, AL, border=NOB)

    # ── 块三（按月）──
    section(ws, '三、成 本（按月 · 你原表按付款对象 果然鲜/伊利/新顺 排，改成按月；按对象看请到【支出月度汇总】【应付账款明细】）',
            'A', 'L', R_SEC3)
    grp(R_G3, R_H3, 'A', 'A', '月份/总支出')
    grp(R_G3, R_G3, 'B', 'K', '成本')
    grp(R_G3, R_H3, 'L', 'L', '小计')
    header(ws, R_H3, 'B', ['所有费用成本\n（果然鲜类）', '装卸费成本', '果然鲜福利', '果然鲜商务', '果然鲜送礼', '筐子成本',
                           '发泡网成本', '包装物料成本', '采购果品', '赔付负数'], 32)
    ws.row_dimensions[R_G3].height = 22
    for rr in (R_T3, R_B3):
        cell(ws, f'A{rr}', '合　计', F_TOT, FL_TOT, ACN)
        for c in 'BCDEFGHIJKL':
            cell(ws, f'{c}{rr}', f'=ROUND(SUM({c}{R_M0}:{c}{R_REST}),2)' if rr == R_T3 else f'={c}{R_T3}',
                 F_TOTN, FL_TOT, AR, MONEY)
    for i, r in enumerate(range(R_M0, R_REST + 1)):
        if r < R_REST:
            cell(ws, f'A{r}', f'=DATE(YEAR($T$7),MONTH($T$7)+{i},1)', F_AUTOB, al=ACN, fmt=MONTH)
            ws[f'T{r}'] = f'=MAX($A{r},$H$2)'
            ws[f'U{r}'] = f'=MIN(EOMONTH($A{r},0),$J$2)'
        else:
            cell(ws, f'A{r}', '其余月份', F_AUTOB, al=ACN)
            ws[f'T{r}'] = '=MAX(DATE(YEAR($T$7),MONTH($T$7)+12,1),$H$2)'
            ws[f'U{r}'] = '=$J$2'
        for c in 'TU':
            ws[f'{c}{r}'].font = F_HELP
            ws[f'{c}{r}'].number_format = DATE
        lo, hi = f'$T{r}', f'$U{r}'
        rp = _period_crit(RJ('B'), lo, hi)
        wp = _period_crit(WL('C'), lo, hi)
        mat = lambda col: (f'IF({lo}>{hi},0,SUMIFS({MAT(col)},{MAT("W")},">="&DATE(YEAR({lo}),MONTH({lo}),1),'
                           f'{MAT("W")},"<="&{hi}))')
        f = {
            'B': (f'SUMIFS({RJ("I")},{RJ("AG")},1,{rp})+SUMIFS({WL("K")},{WL("A")},"03费用应计",{WL("F")},"<>装卸费",{wp})'
                  f'-D{r}-E{r}-F{r}'),
            'C': (f'SUMIFS({WL("K")},{WL("A")},"02装筐费",{wp})+SUMIFS({WL("K")},{WL("A")},"03费用应计",{WL("F")},"装卸费",{wp})'
                  f'+IF($T$4=0,SUMIFS({RJ("I")},{RJ("F")},"装卸费支出",{rp}),0)'),
            'D': f'SUMIFS({RJ("I")},{RJ("AG")},1,{RJ("F")},"福利费",{rp})',
            'E': f'SUMIFS({RJ("I")},{RJ("AG")},1,{RJ("F")},"业务招待费",{RJ("G")},"<>送礼",{rp})',
            'F': f'SUMIFS({RJ("I")},{RJ("AG")},1,{RJ("G")},"送礼",{rp})',
            'G': f'SUMIFS($AD${KZ0}:$AD${KZ1},$AC${KZ0}:$AC${KZ1},">="&{lo},$AC${KZ0}:$AC${KZ1},"<="&{hi})',
            'H': mat('Y'),
            'I': mat('AA'),
            'J': f'SUMIFS({WL("K")},{WL("A")},"01果品",{WL("E")},"应付",{wp})',
            'K': f'SUMIFS({RJ("I")},{RJ("G")},"赔付出库",{rp})',
        }
        for c, ff in f.items():
            cell(ws, f'{c}{r}', f'=IF({lo}>{hi},0,ROUND({ff},2))', al=AR, fmt=MONEY)
        cell(ws, f'L{r}', f'=ROUND(SUM(B{r}:K{r}),2)', F_AUTOB, al=AR, fmt=MONEY)

    # ── 块四（利润）──
    section(ws, '四、利 润（全部按上面三块自动算）', 'A', 'L', R_SEC4)
    grp(R_L4, R_TAG4, 'A', 'A', '利润')
    labels = ['公司利润', '销售果品利润', '发泡网利润', '筐子利润', '包装物料利润', '装卸利润', '营业外利润']
    tags = ['公司收', '个人收', '个人收', '公司收', '个人收', '公司收', '个人收']
    header(ws, R_L4, 'B', labels, 24)
    for i, t in enumerate(tags):
        cell(ws, f'{L(2 + i)}{R_TAG4}', t, F_LBL, FL_LBL, AC)
    grp(R_L4, R_TAG4, 'I', 'I', '合计')
    T1, T2, T3 = R_T1, R_T2, R_T3
    amt = {
        'B': f'=ROUND(D{T1}+E{T1}+K{T1}-(B{T3}+D{T3}+E{T3}+F{T3}),2)',
        'C': f'=ROUND(H{T2}-J{T3},2)',
        'D': f'=ROUND(H{T1}-H{T3},2)',
        'E': f'=ROUND(G{T1}+B{T2}+C{T2}-G{T3},2)',
        'F': f'=ROUND(I{T1}-I{T3},2)',
        'G': f'=ROUND(F{T1}-C{T3},2)',
        'H': f'=ROUND(D{T2}+E{T2}+F{T2}+N(G{T2})+I{T2}+J{T2}-K{T3},2)',
    }
    how = {
        'B': '仓储费＋周转费＋优惠 －（所有费用成本＋福利＋商务＋送礼）',
        'C': '果然鲜销售果品款 － 采购果品',
        'D': '发泡网收入 － 发泡网成本',
        'E': '筐子销售＋筐子/托盘租赁费 － 筐子成本',
        'F': '包装物料收入 － 包装物料成本',
        'G': '装卸费收入 － 装卸费成本',
        'H': '扎伤果款＋垃圾费＋盘盈＋利息＋其他 － 赔付',
    }
    cell(ws, f'A{R_AMT4}', '金额', F_TOT, FL_TOT, ACN)
    for c, f in amt.items():
        cell(ws, f'{c}{R_AMT4}', f, F_TOTN, FL_AUTO, AR, MONEY)
    cell(ws, f'I{R_AMT4}', f'=ROUND(SUM(B{R_AMT4}:H{R_AMT4}),2)', F_TOTN, FL_TOT, AR, MONEY)
    cell(ws, f'A{R_HOW4}', '怎么算', F_LBL, FL_LBL, AC)
    for c, t in how.items():
        cell(ws, f'{c}{R_HOW4}', t, F_NOTE, FL_NONE, AC)
    cell(ws, f'I{R_HOW4}', '七条线合计', F_NOTE, FL_NONE, AC)
    ws.row_dimensions[R_HOW4].height = 48
    ws.row_dimensions[R_AMT4].height = 22
    cell(ws, f'J{R_L4}', '帮我统计个人拿多少利润', F_LBL, FL_LBL, AC)
    cell(ws, f'K{R_L4}', f'=ROUND(C{R_AMT4}+D{R_AMT4}+F{R_AMT4}+H{R_AMT4},2)', F_TOTN, FL_AUTO, AR, MONEY)
    cell(ws, f'J{R_TAG4}', '公司拿多少利润', F_LBL, FL_LBL, AC)
    cell(ws, f'K{R_TAG4}', f'=ROUND(B{R_AMT4}+E{R_AMT4}+G{R_AMT4},2)', F_TOTN, FL_AUTO, AR, MONEY)
    cell(ws, f'L{R_L4}', '个人合计', F_NOTE, FL_NONE, AL, border=NOB)
    cell(ws, f'L{R_TAG4}', '公司合计', F_NOTE, FL_NONE, AL, border=NOB)

    # 与利润表对账
    net, ps, pe = _find_profit_cells(wb)
    r = R_CHK4
    cell(ws, f'A{r}', '与利润表对账', F_LBL, FL_LBL, AC)
    cell(ws, f'B{r}', '本表利润合计', F_LBL, FL_LBL, AC)
    cell(ws, f'C{r}', f'=I{R_AMT4}', F_TOTN, FL_AUTO, AR, MONEY)
    cell(ws, f'D{r}', '利润表净利润', F_LBL, FL_LBL, AC)
    cell(ws, f'E{r}', f'=N(利润表!${net[0]}${net[1:]})', F_TOTN, FL_AUTO, AR, MONEY)
    cell(ws, f'F{r}', '差额', F_LBL, FL_LBL, AC)
    cell(ws, f'G{r}', f'=ROUND(C{r}-E{r},2)', F_TOTN, FL_AUTO, AR, MONEY)
    ws.merge_cells(f'H{r}:L{r}')
    pS, pE = f'利润表!${ps[0]}${ps[1:]}', f'利润表!${pe[0]}${pe[1:]}'
    cell(ws, f'H{r}', f'=IF(OR(N({pS})<>N($H$2),N({pE})<>N($J$2)),"⚠ 利润表期间是 "&TEXT({pS},"yyyy/m/d")&"～"'
                      f'&TEXT({pE},"yyyy/m/d")&"，跟本表不一样：要比，先把利润表第 3 行日期改成跟本表一样",'
                      f'"差额主要来自：仓储/周转/装卸费没进总账；物料、筐子成本没结转；家用进了费用；本表按往来发生额、利润表按凭证")',
         F_NOTE, FL_NONE, AL, border=NOB)
    for c in 'IJKL':
        ws[f'{c}{r}'].border = NOB
    ws.row_dimensions[r].height = 36
    r = R_SELF4
    ws.merge_cells(f'A{r}:L{r}')
    d = f'L{R_T1}+K{R_T2}-L{R_T3}-I{R_AMT4}'
    cell(ws, f'A{r}', f'=IF(ROUND({d},2)=0,"✔ 本表自洽：块一收入＋块二其他收入－块三成本 ＝ 块四七条线利润合计",'
                      f'"✘ 本表收支和利润差 "&TEXT({d},"#,##0.00"))', F_CHK, FL_NONE, AL, border=NOB)

    # 说明 + 要问的问题
    notes = [
        '说明（取不到的、有前提的格子）',
        '1. 盘盈款：系统里没有录盘盈的地方，先写「待定」。实物盘盈不走资金，以后可以在【对账补充行】补一行或告诉我们怎么录。',
        '2. 扎伤梨款 / 扎伤苹果款：取【资金日记账】收入行「费用项目」填「次果梨款／扎伤梨款」「次果苹果款／扎伤苹果款」的金额；'
        '现在没有这类记录，所以是 0。《01》成品出库里的「苹果扎伤」没有单价，算不出钱。垃圾费同理（费用项目＝垃圾费的收入）。',
        '3. 仓储费 / 周转费 / 装卸费：取【原料入库计价】（往来业务明细「03计价」段）。计价日期（M 列）现在全空：没填起止日时'
        '算进来，按月看就算不进去，所以块三、块四没法按月看「公司利润」「装卸利润」。请在【原料入库计价】填计价日期。',
        '4. 所有费用成本（果然鲜类）＝资金日记账里往来单位＝果然鲜、科目是成本/费用类的支出（含人工工资），减去右边单列的'
        '福利、商务、送礼；再加【成本费用登记表】登记的（装卸费除外）。还贷款、家用、往来付款不算。',
        '5. 装卸费成本：《02》装筐费单价设了以后取往来业务明细「02装筐费」（应付口径）；没设之前先按日记账「装卸费支出」'
        '（付给装卸队的钱）算。《02》写「依力」、日记账写「伊利」，要统一成一个名字，不然应付和付款冲不掉。',
        '6. 筐子成本＝卖出去的筐 × 这种筐的采购均价（往来业务明细 02筐子 采购金额÷采购数量，不含期初）。'
        '单价为 0、无需收款的销售出库（如荆建君 10/3 损耗 1 个）不算卖筐，不计成本，所以这里的筐数可能比《02》筐子销售汇总少。',
        '7. 发泡网成本、包装物料成本：取《02》【物料成本接口】（月末一次加权平均，含期初）→【对接源_02物料】W:AB，按整月算；'
        '《02》没更新链接时是 0。',
        '8. 采购果品＝往来业务明细「01果品」应付（成品出库填的采购单价＋原料出库「公司购买」）。赔付取日记账费用项目＝赔付出库，'
        '按正数算进成本（表头保留你原来的「赔付负数」）。',
        '9. 块一 B、C、J 三列你原表表头是空的，先空着；「公司收／个人收」照你的分法拆，公司账里这些收入都记成公司收入，'
        '所以跟利润表口径不一样（见上面「与利润表对账」）。本表收入按应收发生额算，不是按收到的钱。',
        '要问你的 6 个问题',
        '① 块一 B、C、J 三列空表头想放什么？原表块一里「果然鲜」那一行指什么（公司自己卖果？）',
        '② 「所有费用成本（果然鲜类）」包不包括福利、商务、送礼？人工工资算不算公司利润的成本？（现在：不含福利/商务/送礼，含人工工资）',
        '③ 扎伤梨款、扎伤苹果款、垃圾费、盘盈款、赔付，分别在哪里录？是不是公司的钱？',
        '④ 「公司利润」＝仓储费＋周转费－公司费用，装卸单独算，对吗？（现在就是这么算的）',
        '⑤ 「个人收」的几条线（果品、发泡网、包装物料、营业外）在公司账里要不要剔除？',
        '⑥ 仓储/周转/装卸费按哪个日期算月份？（请填【原料入库计价】M 列计价日期，或者定个默认规则，比如按入库月份）',
    ]
    r = R_NOTE
    for t in notes:
        ws.merge_cells(f'A{r}:L{r}')
        if t.startswith('说明') or t.startswith('要问'):
            cell(ws, f'A{r}', t, F_SEC, FL_SEC, AL)
            ws.row_dimensions[r].height = 20
        else:
            cell(ws, f'A{r}', t, Font(name='宋体', size=10, color='1F3864'), FL_NONE, AL, border=NOB)
            # A:L 合并约 166 字符宽，宋体 10 号一行放得下 75 个左右汉字；按字数给够行高（每行 15）
            ws.row_dimensions[r].height = max(30, 15 * -(-len(t) // 70) + 4)
        r += 1

    widths(ws, {'A': 16, 'B': 15, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 14, 'I': 13, 'J': 14,
                'K': 15, 'L': 14, 'S': 10, 'T': 11, 'U': 11, 'V': 12})
    _freeze(ws, 'B3')
    comments = _put_comments(ws, keep, [R_L4, R_TAG4])
    return {'rows': (R_D1, R_E1, R_D2, R_E2, R_M0, R_REST, R_AMT4, R_CHK4), 'profit_cells': (net, ps, pe),
            'comments': comments}


def _grp(ws, r1, r2, c1, c2, text, fill, font):
    ref = f'{c1}{r1}:{c2}{r2}'
    if ref != f'{c1}{r1}:{c1}{r1}':
        ws.merge_cells(ref)
    cell(ws, f'{c1}{r1}', text, font, fill, AC)
    from openpyxl.utils import column_index_from_string as CI
    for rr in range(r1, r2 + 1):
        for cc in range(CI(c1), CI(c2) + 1):
            ws.cell(rr, cc).border = BOX
            if (rr, cc) != (r1, CI(c1)):
                ws.cell(rr, cc).fill = fill


# ───────────────────────────── ④ 主页、使用说明 ─────────────────────────────
def home(wb):
    """主页：改两条旧说明；在最后面另起一块「本轮（10/6）新增/改动 · 报表」（03A 的「· 账务」块在它上面）"""
    ws = wb['主页']
    for r in range(1, ws.max_row + 1):
        b = ws[f'B{r}'].value
        if isinstance(b, str) and "'对接源_01库存'!A1" in b:
            ws[f'C{r}'] = '跨文件取《01》库存等级接口：结余数量、平均重量、在库成品（含损耗）、成品出库总重、损耗'
        if isinstance(b, str) and "'库存价值与欠款比对'!A1" in b:
            ws[f'C{r}'] = ('★上半段按客户比对「原料库存价值＋未发成品价值 ↔ 应收款 ↔ 差额」（200 家）；'
                           '下半段按等级填单价、成品单价、单箱净重')
        if isinstance(ws[f'A{r}'].value, str) and '· 报 表' in ws[f'A{r}'].value:
            return r                                  # 已经加过
    r0 = ws.max_row + 2
    ws.merge_cells(f'A{r0}:D{r0}')
    ws[f'A{r0}'] = '本 轮（10/6）新 增 / 改 动 · 报 表'
    ws[f'A{r0}']._style = copy.copy(ws['A42']._style)
    for c in 'ABCD':
        ws[f'{c}{r0 + 1}'] = ws[f'{c}43'].value
        ws[f'{c}{r0 + 1}']._style = copy.copy(ws[f'{c}43']._style)
    items = [
        ('果然鲜总表', '★你建的总表改成全自动：按客户的收入、其他收入、按月成本、各条线利润（公司收／个人收），带时段和与利润表对账',
         '自动'),
        ('库存价值与欠款比对', '下半段加「在库成品吨位毛重、估算净重、成品单价（手填）、未发成品价值」；上半段扩到 200 家，'
                               '差额＝库存价值＋未发成品价值－应收', '录入+自动'),
        ('对接源_01库存', '加 L/M/N 三列：在库成品（含损耗）、成品出库总重、损耗（取《01》库存等级接口，不随日期筛选）', '自动'),
    ]
    for i, (nm, txt, attr) in enumerate(items):
        r = r0 + 2 + i
        ws[f'A{r}'] = i + 1
        ws[f'B{r}'] = f'=HYPERLINK("#\'{nm}\'!A1","{nm}")'
        ws[f'C{r}'] = txt
        ws[f'D{r}'] = attr
        for c in 'ABCD':
            ws[f'{c}{r}']._style = copy.copy(ws[f'{c}55']._style)
    return r0


def manual(wb):
    ws = wb['使用说明']
    if any(isinstance(ws[f'A{r}'].value, str) and '第十版 · 报表' in ws[f'A{r}'].value for r in range(1, ws.max_row + 1)):
        return
    last = max(r for r in range(1, ws.max_row + 1) if any(ws.cell(r, c).value is not None for c in range(1, 9)))
    hs = copy.copy(ws['A123']._style)
    ls = copy.copy(ws['A124']._style)
    lines = [
        ('★ 本 次 更 新（2026-10-06 · 第十版 · 报表）', hs, 19.5),
        ('1、【库存价值与欠款比对】上半段扩到 200 家、加「结余成品吨位净重」「在库成品毛重」「未发成品价值」，'
         '差额＝库存价值＋未发成品价值－应收款；下半段整体挪到第 208 行以后，单价后面加「在库成品吨位毛重（自动）」'
         '「估算在库净重」「成品单价（手填）」「未发成品价值」，你手填的单价、单箱净重都原样搬过去了（P 列是单箱净重）。', ls, 45),
        ('2、【果然鲜总表】改成全自动：第 2 行选年份/起止日；块一按客户列收入，块二其他收入，块三按月列成本，'
         '块四各条线利润分公司收／个人收，最下面写了取不到的格子和要问你的问题。', ls, 45),
    ]
    r = last + 2
    for t, st, h in lines:
        ws.merge_cells(f'A{r}:H{r}')
        ws[f'A{r}'] = t
        ws[f'A{r}']._style = st
        ws.row_dimensions[r].height = h
        r += 1


def apply(wb):
    for name in (SHEET_LV, SRC, ZB, '主页', '使用说明', '往来业务明细', '资金日记账', '利润表', '对接源_02物料'):
        if name not in wb.sheetnames:
            raise SystemExit(f'《03》里没有【{name}】')
    source_01(wb)
    info = stock_compare(wb)
    info.update(summary(wb))
    info['home_row'] = home(wb)
    manual(wb)
    return info
