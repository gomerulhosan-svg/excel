# -*- coding: utf-8 -*-
"""《03》主页入口、使用说明补这一轮；《04》果然鲜销售对账单加「购买人」列。"""
import copy
from openpyxl.worksheet.hyperlink import Hyperlink
from common0928 import *
from xlsx_util import insert_column

NEW_03 = [
    ('往来对账单', '★原版式对账单：收付款、定金押金、借款、物料、筐子、仓储装卸周转费全在一张，逐笔带期末余额'),
    ('往来对账单明细', '每家单位一行：期初、本期金额、已付款、期末余额（拆应收/应付），带合计；【库存价值与欠款比对】取这里的期末余额'),
    ('应收账款对账单', '同一版式，只看应收方向；每行期末余额'),
    ('应付账款对账单', '同一版式，只看应付方向；每行期末余额'),
    ('支出月度汇总', '往来单位 × 12 个月统计支出（不含账户互转），带小计、合计和核对'),
    ('费用月度支出汇总', '果然鲜公司费用：费用项目 × 12 个月'),
    ('家用月度支出汇总', '家用：按往来单位（家小赫、家小满…）× 月、按费用项目 × 月'),
    ('家用月度明细汇总', '家用：选年份/月份，左边按往来单位、费用项目看收入支出，右边逐笔明细'),
    ('国外费用汇总', '业务类型「国外」：全年给国外转出多少，按往来单位汇总＋逐笔明细'),
    ('借款汇总', '借入、借出分开：总借款、总还款、余额，以及全部借款结余合计；可补记账前的老借款'),
    ('往来业务明细', '所有往来业务逐笔底表（对账单、汇总表都从这里取数），勿手工改'),
]

DESC_03 = {
    '资金与往来余额表': '改回原样式：资金账户结存（带合计、与账面货币资金的差额核对）＋往来单位余额（带合计）；选年份/月份/起止日',
    '应收应付汇总表': '每家单位应收、应付分方向＋净往来，带时段和合计',
    '收入月度汇总': '★往来单位 × 12 个月统计收入，选年份/起止日，带合计和核对',
    '应付账款明细': '只列有应付业务的单位：期初、本期发生、已付、余额，带合计',
}


def home_03(wb):
    ws = wb['主页']
    # 旧链接修正：费用月度汇总 → 费用月度明细汇总
    for r in range(7, 40):
        b = ws[f'B{r}']
        if b.value == '费用月度汇总':
            b.value = '费用月度明细汇总'
            b.hyperlink = Hyperlink(ref=f'B{r}', location="'费用月度明细汇总'!A1", display='费用月度明细汇总')
            ws[f'C{r}'] = '★公司费用（往来单位＝果然鲜）按费用项目汇总＋逐笔明细，选年份/月份/起止日，带合计'
        elif b.value == '往来对账单':
            ws[f'C{r}'] = NEW_03[0][1]
        elif b.value in DESC_03:
            ws[f'C{r}'] = DESC_03[b.value]
    # 「本轮新增」段：保留原两行，后面接这一轮的
    ws['A42'] = '本 轮 新 增 / 改 动'
    src = [ws[f'{c}45'] for c in 'ABCD']
    start = 46
    shown = [x for x in NEW_03 if x[0] != '往来对账单']
    for i, (name, desc) in enumerate(shown):
        r = start + i
        for c, s in zip('ABCD', src):
            d = ws[f'{c}{r}']
            d._style = copy.copy(s._style)
        ws[f'A{r}'] = 3 + i
        ws[f'B{r}'] = f'=HYPERLINK("#\'{name}\'!A1","{name}")'
        ws[f'C{r}'] = desc
        ws[f'D{r}'] = '录入+自动' if name in ('借款汇总',) else '自动'
        ws[f'C{r}'].alignment = copy.copy(ws[f'C{r}'].alignment)
        ws.row_dimensions[r].height = 30 if len(desc) > 38 else ws.row_dimensions[45].height


def manual_03(wb):
    ws = wb['使用说明']
    hdr, item = ws['A116'], ws['A117']
    r = ws.max_row + 2
    lines = [
        '★ 本 次 更 新（2026-09-28 · 第九版）',
        '1、【往来对账单】改回原来「新疆果然鲜仓储有限公司 对账单」版式，而且不再分应收应付：收款、付款、定金押金的收和退、退款、借款、'
        '包装物料、筐子、仓储/装卸/周转费、公司购买客户果品，只要跟这个人有关的全在一张单上，按日期逐笔排，每行带「期末余额」。'
        '金额＝我方应收的（负数＝我方应付），已付款＝对方付给我方的（负数＝我方付出去的），期末余额正数＝对方欠我方。',
        '2、【应收账款对账单】【应付账款对账单】也换成同一版式，「余额变动」改成每行的「期末余额」。定金押金在会计上是应付，'
        '所以只看应收对账单时看不到，要看全口径请用【往来对账单】。',
        '3、新增【往来对账单明细】：每家单位一行，期初／本期金额／已付款／期末余额，带合计；【库存价值与欠款比对】的应收款余额改取这里的期末余额。',
        '4、【资金与往来余额表】改回原样式（资金账户结存＋往来单位余额），保留合计和差额核对，可以直接选年份、月份。',
        '5、新增【支出月度汇总】【费用月度支出汇总】【家用月度支出汇总】【家用月度明细汇总】【国外费用汇总】【借款汇总】；'
        '所有统计表第 2 行都有「年份／起始日／截止日」（查月的表还有「月份」），都带合计。',
        '6、删除【月度资金明细】；【资金日记账】里只给它和旧汇总用的 AB～AE 辅助列一并清掉。',
        '7、【往来业务明细】是逐笔底表，每行只是这一笔的增减（列名改成「本笔增减」）；每家单位的期末余额看三张对账单和【往来对账单明细】。',
        '8、【对账补充行】能用的行数从 100 行扩到 300 行（原来第 104 行以后填了也不进对账单）。记账以前收的定金、欠的款，'
        '可以在这里补一行期初（例如麦麦提有退定金、没有收定金的记录）。',
    ]
    for i, t in enumerate(lines):
        rr = r + i
        ws.merge_cells(f'A{rr}:H{rr}')
        ws[f'A{rr}'] = t
        ws[f'A{rr}']._style = copy.copy((hdr if i == 0 else item)._style)
        ws.row_dimensions[rr].height = 19.5 if i == 0 else 45


# ─────────────────────────────── 《04》 ───────────────────────────────
def buyer_04(wb):
    name = '果然鲜销售对账单'
    ws = wb[name]
    insert_column(wb, name, 3, 1210, 20, width=12)
    ws['C5'] = '购 买 人'
    ws['C5']._style = copy.copy(ws['B5']._style)
    for r in range(6, 206):
        ws[f'C{r}'] = (f'=IF($T{r}="","",IF(INDEX(对接源_01水果!$AC$4:$AC$1203,$T{r})=0,"",'
                       f'INDEX(对接源_01水果!$AC$4:$AC$1203,$T{r})))')
        ws[f'C{r}']._style = copy.copy(ws[f'B{r}']._style)
    ws['C206']._style = copy.copy(ws['B206']._style)
    ws.freeze_panes = 'D6'
    # insert_column 只挪单列宽度，原来 G:J、Q:R 两段合并宽度要手工补回（右移一格后是 H:K、R:S）
    for c in 'HIJK':
        ws.column_dimensions[c].width = 11
    for c in 'RS':
        ws.column_dimensions[c].width = 14
    ws.column_dimensions['G'].width = 10          # 页脚「合计金额」的数落在这一列，原来 7 宽会显示 ###
    if ws['A208'].value and not any(ws[f'{c}208'].value for c in 'BCDE'):
        ws.merge_cells('A208:E208')
    ws['A3'] = ('在 B2 选择送礼、销售、商务或福利，按销售类型及起止日期提取；日期留空不限。C 列「购买人」取《01》成品出库明细的'
                '购买方（没填购买方就是客户/领用方）。零金额业务也显示，金额沿用来源记录。')
    h = wb['主页']
    for r in range(1, 40):
        if isinstance(h[f'B{r}'].value, str) and '果然鲜销售对账单' in h[f'B{r}'].value:
            h[f'C{r}'] = '★按销售类型（送礼/销售/商务/福利）出：逐笔列购买人、卖了什么、多少钱、运费多少、收了多少、还欠多少'


# 新表跟前面老表一样：横向 A4、按宽度缩到一页宽、每页重复表头；数据行 18 高
PAGE_03 = {
    '往来业务明细': '1:3', '往来对账单明细': '1:5', '应付账款明细': '1:5', '应收应付汇总表': '1:5',
    '资金与往来余额表': '1:2', '收入月度汇总': '1:6', '支出月度汇总': '1:6', '费用月度支出汇总': '1:6',
    '家用月度支出汇总': '1:5', '家用月度明细汇总': '1:6', '费用月度明细汇总': '1:6', '国外费用汇总': '1:6',
    '借款汇总': '1:4',
}


def finish_03(wb):
    for name, rows in PAGE_03.items():
        ws = wb[name]
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = 9
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = rows
        for r in range(4, ws.max_row + 1):
            if ws.row_dimensions[r].height is None:
                ws.row_dimensions[r].height = 18
