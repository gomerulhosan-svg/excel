# -*- coding: utf-8 -*-
"""10/6 这一轮（A050 1006）《01 水果进销存》的改动。底稿：参考/1006上传/01_水果进销存台账模板.xlsx。

用法：wb = openpyxl.load_workbook(底稿)；apply(wb)；wb.save(成品)；post(成品)。
对底稿重复跑，结果一样（所有公式、样式都按固定规则重写，不依赖上一次的结果）。

这一轮改了什么（对应 A050 1006 实施规格第 2 节、接口 K1 / K5 / K8）：

① 修「中间插行就漏最后一笔」：【_自动清单】整张重写，所有逐行镜像都改成按位置取
   INDEX(明细!$F$4:$F$5003,ROW()-3)，以后在明细中间插行也不会错位；紧凑清单都加了「超过个数就不查」护栏；
   25,600 个 INDIRECT（果然鲜采购/销售明细、_自动清单 的累计列）全部换成普通累计＋INDEX；
   原料出库明细 K103（用户插的那行）补上总重公式。
② 扩容：原料入库明细 4～5003（5,000 行，公式/格式/下拉/筛选一起扩）；对账明细接口 4～6003（6,000 行，K1）；
   池子分段：原料入库 1～5000、原料出库 5001～8001、成品出库 8002～11001、其他代存 11002～11801。
③ 库存结余：H 后插 I「平均重量KG/件」（全部入库总重÷入库数量，不随日期）；出库总重＝出库数量×平均重量；
   结余总重＝入库总重－出库总重；成品出库总重后加 P「在库成品吨位（含损耗）KG」＝加工类出库（出库加工、二次加工）
   数量×平均重量－成品出库总重；Q「损耗KG」（方案 A：原料没出完留空、备注「原料未出完」；出完后
   ＝入库总重－直接出库数量×平均重量－成品出库总重，报损算损耗）；R「损耗率」＝损耗÷加工出库重量。负数标红并备注。
   组合清单把成品出库明细的组合也收进来（扎伤、货主＝果然鲜），排在入库、出库组合后面。
   合计行补入库数量，删掉对「库存状态」文字求和的那格。库存总结余 H、实盘库存盘点 H 的出库重量改成同口径。
④ 库存等级接口（K8）：A～H 含义不变，补 I 平均重量，新增 J 在库成品（含损耗）、K 成品出库总重、L 损耗、
   M 加工出库重量；整张接口按全部日期算，不跟【库存结余】的日期筛选走；计数格挪到 R1:S1。
⑤ 原料出库明细加 U 采购单价（元/件）、V 采购金额（应付）、W 付款状态、X 购买方/去向。
⑥ 果然鲜采购明细并入原料出库「公司购买」的行（先成品出库、后原料出库，加「来源」列），去掉 INDIRECT。
⑦ 对账明细接口：原料出库「公司购买」行也输出 T/U/V/W/X/AC/AF；原料出库 J 列重量没填时用 数量×平均重量。
⑧ 财务取数接口第三块：公司购买果品款（应付/已付）加上原料出库「公司购买」；506 行笔误统一。
⑨ 新表【公司购买接口】（K5）：原料出库「公司购买」逐笔，不带日期筛选，供《03》取数。
⑩ 果然鲜销售汇总并进果然鲜销售明细「三、按购买方汇总」（第 478 行起），删掉原表；二段改成扫全部 400 行明细，
   「品种数」改成真的品种数，J 列宽修好。
⑪ 主页加新表入口、改链接；使用说明补这一轮要点。
   公司购买的金额在 _自动清单 CX 列现算（出库数量×采购单价），对账明细接口 / 公司购买接口 / 果然鲜采购明细 /
   财务取数接口都取这一列，用户在中间插行、新行没带 V 列公式也不会漏金额。

post(path)：核对定义名称（删表后 localSheetId）、没有外链；把 openpyxl 写成 &#xxxx; 的中文改回 UTF-8，
【对账明细接口】【_自动清单】改成共享公式（WPS/Excel 自己存盘也是这样），文件 9MB → 3.7MB，算出来不变。
"""
import copy
import re
import zipfile

from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font, Alignment

from common0928 import (F_TITLE, F_NOTE, F_HDR, F_IN, F_AUTO, F_AUTOB, F_TOT, F_TOTN, F_HELP, F_LBL,
                        FL_TITLE, FL_HDR, FL_AUTO, FL_IN, FL_TOT, FL_LBL, FL_NOTE, FL_NONE,
                        AC, AL, BOX, NOB, MONEY, QTY, INT, DATE, cell)

# ── 各录入表的数据区（第 4 行起）──────────────────────────────────────────────
IN_S, OUT_S, FIN_S, DAI_S = '原料入库明细', '原料出库明细', '成品出库明细', '其他代存明细'
IN_END, OUT_END, FIN_END, DAI_END = 5003, 3004, 3003, 803        # 原料入库扩到 5,000 行
IN_N, OUT_N, FIN_N, DAI_N = IN_END - 3, OUT_END - 3, FIN_END - 3, DAI_END - 3   # 5000 / 3001 / 3000 / 800

# _自动清单 里「原料入库 → 原料出库 → 成品出库 → 其他代存」首尾相接的池子
P_IN0 = 4
P_OUT0 = P_IN0 + IN_N            # 5004
P_FIN0 = P_OUT0 + OUT_N          # 8005
P_DAI0 = P_FIN0 + FIN_N          # 11005
P_END3 = P_DAI0 - 1              # 11004：组合 / 筐 池子（入库＋出库＋成品）
P_END4 = P_DAI0 + DAI_N - 1      # 11804：对账命中池子（再加其他代存）
# 对账明细接口 Z 列（池子里的第几个）分段：≤5000 入库，≤8001 出库，≤11001 成品，其余代存
T_IN, T_OUT, T_FIN = IN_N, IN_N + OUT_N, IN_N + OUT_N + FIN_N

DZ_END = 6003                    # 对账明细接口 4～6003（K1）
KC_END = 403                     # 库存结余 / 库存等级接口 4～403
GM_END = 403                     # 公司购买接口 4～403（K5）

MONEY_FMT = '#,##0.00;[Red]\\-#,##0.00;\\-'
PCT_FMT = '0.0%;[Red]\\-0.0%;\\-'


def R(sheet, col, end, start=4):
    return f'{sheet}!${col}${start}:${col}${end}'


def IX(sheet, col, end, p):
    """INDEX(明细!$X$4:$X$end, p)"""
    return f'INDEX({R(sheet, col, end)},{p})'


def V(sheet, col, end, p):
    """取一格：空就给 ""（跟原来 IF(x=0,"",x) 的写法一样）"""
    x = IX(sheet, col, end, p)
    return f'IF({x}=0,"",{x})'


def NV(sheet, col, end, p):
    """取一格数字：0 就给 \"\" """
    x = IX(sheet, col, end, p)
    return f'IF(N({x})=0,"",N({x}))'


def style_from(dst, src):
    dst._style = copy.copy(src._style)


def date_num(x):
    """日期格（可能是文字 2026-09-01）转成日期序号，跟原表 CA/CJ 列的写法一致"""
    return (f'IF(ISNUMBER({x}),{x},IFERROR(DATE(VALUE(LEFT({x},4)),VALUE(MID({x},6,2)),'
            f'VALUE(MID({x},9,2))),0))')


# ════════════════════════════════════════════════════════════════════════════
# ① ② 录入表：原料出库明细加 4 列、补 K103；原料入库明细扩到 5,000 行
# ════════════════════════════════════════════════════════════════════════════
def out_sheet(wb):
    ws = wb[OUT_S]
    fin = wb[FIN_S]
    # K103（用户插的那行）没带公式，整列统一补齐
    for r in range(4, OUT_END + 1):
        if ws[f'K{r}'].value is None:
            ws[f'K{r}'] = f'=IF(OR($H{r}="",$J{r}=""),"",ROUND($H{r}*$J{r},2))'
            style_from(ws[f'K{r}'], ws['K4'])
    # U～X 四列
    heads = {'U': '采购单价\n（元/件）', 'V': '采购金额\n（应付）', 'W': '付款状态', 'X': '购买方/去向'}
    for c, t in heads.items():
        ws[f'{c}3'] = t
        style_from(ws[f'{c}3'], fin['T3'])          # 橙色「采购」表头，跟成品出库明细一致
    for r in range(4, OUT_END + 1):
        u, v, w, x = ws[f'U{r}'], ws[f'V{r}'], ws[f'W{r}'], ws[f'X{r}']
        style_from(u, fin['T4'])                      # 浅蓝录入
        u.number_format = '#,##0.00;[Red]\\-#,##0.00;\\-'
        v.value = f'=IF(OR($B{r}="",N($U{r})=0),"",ROUND(N($H{r})*N($U{r}),2))'
        style_from(v, ws['K4'])                       # 浅绿公式
        v.number_format = MONEY_FMT
        style_from(w, fin['X4'])
        w.number_format = 'General'
        style_from(x, fin['X4'])
        x.number_format = 'General'
    dv = DataValidation(type='list', formula1='"已付清,部分付款,未付款,无需付款"', allow_blank=True)
    dv.add(f'W4:W{OUT_END}')
    ws.add_data_validation(dv)
    dv = DataValidation(type='list', formula1='购买方表', allow_blank=True, showErrorMessage=False)
    dv.add(f'X4:X{OUT_END}')
    ws.add_data_validation(dv)
    for c, w in {'U': 11, 'V': 12, 'W': 10, 'X': 14}.items():
        ws.column_dimensions[c].width = w
    for m in [str(m) for m in ws.merged_cells.ranges]:
        if m in ('A1:T1', 'A2:T2'):
            ws.unmerge_cells(m)
            ws.merge_cells(m.replace('T', 'X'))
    a2 = ws['A2'].value or ''
    if '同一批货别两头都填' not in a2:
        ws['A2'] = (a2 + '\n★ 右边 U～X 是「公司购买」用的：公司在原料阶段从货主手里买的货（出库原因选「公司购买」），'
                    '采购单价（元/件，按出库数量算）填在这里，金额自动算；直接买货主加工好的成品，才填【成品出库明细】的采购单价'
                    '——同一批货别两头都填。')
    ws.row_dimensions[2].height = 48
    ws.auto_filter.ref = f'A3:X{OUT_END}'


def in_sheet(wb):
    """原料入库明细 4～3004 → 4～5003：公式、格式、下拉、筛选一起扩"""
    ws = wb[IN_S]
    old_end = 3004
    tmpl = {c: ws[f'{c}{old_end}'] for c in [L(i) for i in range(1, 18)]}   # A～Q
    for r in range(old_end + 1, IN_END + 1):
        for c, src in tmpl.items():
            d = ws[f'{c}{r}']
            if src.has_style:
                d._style = copy.copy(src._style)
        ws[f'A{r}'] = f'=IF($B{r}="","",ROW()-3)'
        ws[f'K{r}'] = f'=IF(OR($H{r}="",$J{r}=""),"",ROUND($H{r}*$J{r},2))'
        ws[f'M{r}'] = f'=IF(OR($L{r}="",$H{r}=""),"",$H{r})'
    h = ws.row_dimensions[old_end].height
    if h:
        for r in range(old_end + 1, IN_END + 1):
            ws.row_dimensions[r].height = h
    for dv in ws.data_validations.dataValidation:
        parts = str(dv.sqref).split()
        parts = [re.sub(rf'(\D)({old_end})$', rf'\g<1>{IN_END}', p) for p in parts]
        dv.sqref = type(dv.sqref)(' '.join(parts))
    ws.auto_filter.ref = f'A3:Q{IN_END}'


def widen_refs(wb, skip):
    """全册公式里 原料入库明细!$X$4:$X$3004 → $5003；_自动清单!$S$4:$S$3003 → $5003（库存总结余用的入库货主品种池）"""
    pat = re.compile(r"(原料入库明细'?!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)3004\b")
    pat2 = re.compile(r"(_自动清单!\$S\$4:\$S\$)3003\b")
    n = 0
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('=') and ('原料入库明细!' in v or '_自动清单!$S$' in v):
                    nv = pat2.sub(rf'\g<1>{IN_END}', pat.sub(rf'\g<1>{IN_END}', v))
                    if nv != v:
                        c.value = nv
                        n += 1
    return n


# ════════════════════════════════════════════════════════════════════════════
# ① _自动清单：整张按位置重写
# ════════════════════════════════════════════════════════════════════════════
def auto_list(wb):
    ws = wb['_自动清单']
    hdr_style = copy.copy(ws['A3']._style)
    a2 = ws['A2'].value
    a2_style = copy.copy(ws['A2']._style)
    old_hdr = {c.column_letter: c.value for c in ws[3] if c.value is not None}
    ws._cells.clear()
    ws['A2'] = a2
    ws['A2']._style = a2_style
    hdr = dict(old_hdr)
    hdr.update({'G': None, 'O': None, 'J': '位置', 'X': '位置', 'AH': '位置', 'AT': '位置', 'BF': '位置', 'BN': '位置',
                'CC': '成品日期', 'CG': '位置',
                'CP': '公司购买', 'CQ': '累计', 'CR': '采购明细时段内', 'CS': '累计',
                'CT': '估算出库重量', 'CU': '出库日期', 'CV': '货主', 'CW': '品种',
                'CX': '公司购买金额', 'CY': '付款状态'})
    for c, t in hdr.items():
        if t is None:
            continue
        ws[f'{c}3'] = t
        ws[f'{c}3']._style = copy.copy(hdr_style)

    def put(ref, f):
        ws[ref] = f

    # ---- 每段的取数写法 ------------------------------------------------------
    segs3 = [  # 组合池、筐池：入库、出库、成品
        ('IN', IN_S, IN_END, P_IN0, P_OUT0 - 1, P_IN0 - 1),
        ('OUT', OUT_S, OUT_END, P_OUT0, P_FIN0 - 1, P_OUT0 - 1),
        ('FIN', FIN_S, FIN_END, P_FIN0, P_END3, P_FIN0 - 1),
    ]
    # A～I：货主|品种|等级（库存结余用）
    for name, sh, end, r0, r1, off in segs3:
        p = f'ROW()-{off}'
        F_, D_, E_ = IX(sh, 'F', end, p), IX(sh, 'D', end, p), IX(sh, 'E', end, p)
        for r in range(r0, r1 + 1):
            put(f'A{r}', f'=IF({F_}="","",{F_}&"|"&{D_}&"|"&{E_})')
            put(f'B{r}', f'=IF(A{r}="",0,IF(MATCH(A{r},$A$4:$A${P_END3},0)=ROW()-3,1,0))')
            put(f'C{r}', f'=N(C{r - 1})+B{r}')
            put(f'D{r}', f'=IF($A{r}="","",{F_})')
            put(f'E{r}', f'=IF($A{r}="","",{D_})')
            put(f'F{r}', f'=IF($A{r}="","",{E_})')
            if name != 'FIN':
                put(f'H{r}', f'=IF($A{r}="","",{IX(sh, "G", end, p)})')
                put(f'I{r}', f'=IF($A{r}="","",{IX(sh, "I", end, p)})')
            else:
                put(f'I{r}', f'=IF($A{r}="","",{IX(sh, "K", end, p)})')
    put('A1', f'=MAX($C$4:$C${P_END3})')
    for r in range(4, KC_END + 1):
        put(f'J{r}', f'=IF(ROW()-3>$A$1,"",MATCH(ROW()-3,$C$4:$C${P_END3},0))')
        for c, src in zip('KLMN', 'ADEF'):
            put(f'{c}{r}', f'=IF($J{r}="","",INDEX(${src}$4:${src}${P_END3},$J{r}))')
        for c, src in zip('PQ', 'HI'):
            put(f'{c}{r}', f'=IF($J{r}="","",INDEX(${src}$4:${src}${P_END3},$J{r})&"")')

    # S～W：入库 货主|品种（库存总结余、入库吨位汇总）
    p = 'ROW()-3'
    F_, D_ = IX(IN_S, 'F', IN_END, p), IX(IN_S, 'D', IN_END, p)
    for r in range(4, IN_END + 1):
        put(f'S{r}', f'=IF({F_}="","",{F_}&"|"&{D_})')
        put(f'T{r}', f'=IF(S{r}="",0,IF(MATCH(S{r},$S$4:$S${IN_END},0)=ROW()-3,1,0))')
        put(f'U{r}', f'=N(U{r - 1})+T{r}')
        put(f'V{r}', f'=IF($S{r}="","",{F_})')
        put(f'W{r}', f'=IF($S{r}="","",{D_})')
    put('S1', f'=MAX($U$4:$U${IN_END})')
    compact(ws, 'X', ['Y', 'Z', 'AA'], ['S', 'V', 'W'], 'U', '$S$1', IN_END, 303)

    # AC～AG：成品 货主|品种（成品出库吨位汇总）
    F_, D_ = IX(FIN_S, 'F', FIN_END, p), IX(FIN_S, 'D', FIN_END, p)
    for r in range(4, FIN_END + 1):
        put(f'AC{r}', f'=IF({F_}="","",{F_}&"|"&{D_})')
        put(f'AD{r}', f'=IF(AC{r}="",0,IF(MATCH(AC{r},$AC$4:$AC${FIN_END},0)=ROW()-3,1,0))')
        put(f'AE{r}', f'=N(AE{r - 1})+AD{r}')
        put(f'AF{r}', f'=IF($AC{r}="","",{F_})')
        put(f'AG{r}', f'=IF($AC{r}="","",{D_})')
    put('AC1', f'=MAX($AE$4:$AE${FIN_END})')
    compact(ws, 'AH', ['AI', 'AJ', 'AK'], ['AC', 'AF', 'AG'], 'AE', '$AC$1', FIN_END, 303)

    # AM～AS：其他代存 客户|品种|等级
    G_, E_, F_, I_ = (IX(DAI_S, c, DAI_END, p) for c in 'GEFI')
    for r in range(4, DAI_END + 1):
        put(f'AM{r}', f'=IF({G_}="","",{G_}&"|"&{E_}&"|"&{F_})')
        put(f'AN{r}', f'=IF(AM{r}="",0,IF(MATCH(AM{r},$AM$4:$AM${DAI_END},0)=ROW()-3,1,0))')
        put(f'AO{r}', f'=N(AO{r - 1})+AN{r}')
        put(f'AP{r}', f'=IF($AM{r}="","",{G_})')
        put(f'AQ{r}', f'=IF($AM{r}="","",{E_})')
        put(f'AR{r}', f'=IF($AM{r}="","",{F_})')
        put(f'AS{r}', f'=IF($AM{r}="","",{I_})')
    put('AM1', f'=MAX($AO$4:$AO${DAI_END})')
    compact(ws, 'AT', ['AU', 'AV', 'AW', 'AX', 'AY'], ['AM', 'AP', 'AQ', 'AR', 'AS'], 'AO', '$AM$1', DAI_END, 203)

    # BA～BE：客户|筐子类型（筐子进出接口）：入库、出库按货主＋筐子类型，成品按客户＋筐子来源
    for name, sh, end, r0, r1, off in segs3:
        pp = f'ROW()-{off}'
        if name == 'FIN':
            who, box = IX(sh, 'G', end, pp), IX(sh, 'O', end, pp)
        else:
            who, box = IX(sh, 'F', end, pp), IX(sh, 'L', end, pp)
        for r in range(r0, r1 + 1):
            put(f'BA{r}', f'=IF(OR({box}="",{who}=""),"",{who}&"|"&{box})')
            put(f'BB{r}', f'=IF(BA{r}="",0,IF(MATCH(BA{r},$BA$4:$BA${P_END3},0)=ROW()-3,1,0))')
            put(f'BC{r}', f'=N(BC{r - 1})+BB{r}')
            put(f'BD{r}', f'=IF($BA{r}="","",{who})')
            put(f'BE{r}', f'=IF($BA{r}="","",{box})')
    put('BA1', f'=MAX($BC$4:$BC${P_END3})')
    compact(ws, 'BF', ['BG', 'BH', 'BI'], ['BA', 'BD', 'BE'], 'BC', '$BA$1', P_END3, 303)

    # BK～BO：往来单位（货主 L、成品货主 AJ、筐客户 BH、代存客户 AV）——原料出库「公司购买」的卖方就是货主，已在 L 里
    for r in range(4, 1204):
        if r <= 403:
            s = f'$L{r}'
        elif r <= 703:
            s = f'$AJ{r - 400}'
        elif r <= 1003:
            s = f'$BH{r - 700}'
        else:
            s = f'$AV{r - 1000}'
        put(f'BK{r}', f'=IF({s}="","",{s})')
        put(f'BL{r}', f'=IF(BK{r}="",0,IF(MATCH(BK{r},$BK$4:$BK$1203,0)=ROW()-3,1,0))')
        put(f'BM{r}', f'=N(BM{r - 1})+BL{r}')
    put('BK1', '=MAX($BM$4:$BM$1203)')
    compact(ws, 'BN', ['BO'], ['BK'], 'BM', '$BK$1', 1203, 253)

    # BQ～BS：计量单位（原样，只是区间扩到 5003）
    for r in range(4, 104):
        put(f'BQ{r}', f'=IF(基础资料!$G{r}="","",基础资料!$G{r})')
        put(f'BR{r}', f'=IF(BQ{r}="",0,COUNTIF({R(IN_S, "I", IN_END)},BQ{r}))')
        put(f'BS{r}', f'=IF(BQ{r}="",0,COUNTIF({R(FIN_S, "K", FIN_END)},BQ{r}))')
    put('BQ1', '=IF(MAX($BR$4:$BR$103)=0,"筐",INDEX($BQ$4:$BQ$103,MATCH(MAX($BR$4:$BR$103),$BR$4:$BR$103,0)))')
    put('BR1', '=IF(MAX($BS$4:$BS$103)=0,"箱",INDEX($BQ$4:$BQ$103,MATCH(MAX($BS$4:$BS$103),$BS$4:$BS$103,0)))')

    # BX～BY：对账命中（入库、出库、成品、代存首尾相接）
    for name, sh, end, r0, r1, off in segs3 + [('DAI', DAI_S, DAI_END, P_DAI0, P_END4, P_DAI0 - 1)]:
        key = {'IN': 'F', 'OUT': 'F', 'FIN': 'G', 'DAI': 'G'}[name]
        k = IX(sh, key, end, f'ROW()-{off}')
        for r in range(r0, r1 + 1):
            put(f'BX{r}', f'=IF({k}="",0,1)')
            put(f'BY{r}', f'=N(BY{r - 1})+$BX{r}')
    put('BX1', f'=MAX($BY$4:$BY${P_END4})')

    # CA～CN：成品出库明细逐行（果然鲜销售命中、采购命中、购买方）
    B_ = IX(FIN_S, 'B', FIN_END, p)
    for r in range(4, FIN_END + 1):
        I_, Y_, AB_, T_, V_, H_, G_ = (IX(FIN_S, c, FIN_END, p) for c in ('I', 'Y', 'AB', 'T', 'V', 'H', 'G'))
        put(f'CC{r}', f'=IF({B_}="",0,{date_num(B_)})')
        put(f'CA{r}', f'=IF(AND({B_}<>"",OR(LEFT(TRIM({I_}),3)="果然鲜",IFERROR(VALUE({Y_}),0)<>0,'
                      f'IFERROR(VALUE({AB_}),0)<>0),$CC{r}>=果然鲜销售明细!$X$1,$CC{r}<=果然鲜销售明细!$X$2),1,0)')
        put(f'CB{r}', f'=N(CB{r - 1})+CA{r}')
        put(f'CJ{r}', f'=IF(AND({B_}<>"",OR(IFERROR(VALUE({T_}),0)<>0,IFERROR(VALUE({V_}),0)<>0),'
                      f'$CC{r}>=果然鲜采购明细!$Q$1,$CC{r}<=果然鲜采购明细!$Q$2),1,0)')
        put(f'CK{r}', f'=N(CK{r - 1})+CJ{r}')
        put(f'CD{r}', f'=IF(AND(CA{r}=0,CJ{r}=0),"",IF({H_}<>"",{H_},IF({G_}=0,"",{G_})))')
        put(f'CM{r}', f'=IF(AND({B_}<>"",OR(N({Y_})<>0,N({AB_})<>0)),1,0)')
        put(f'CN{r}', f'=N(CN{r - 1})+CM{r}')
    # 购买方池子后半段：原料出库「公司购买」（在采购明细时段内）的购买方/去向
    for r in range(FIN_END + 1, FIN_END + OUT_N + 1):
        x = IX(OUT_S, 'X', OUT_END, f'ROW()-{FIN_END}')
        put(f'CD{r}', f'=IF(N($CR{r - FIN_N})<>1,"",IF({x}=0,"",{x}))')
    BUY_END = FIN_END + OUT_N   # 6004
    for r in range(4, BUY_END + 1):
        put(f'CE{r}', f'=IF(CD{r}="",0,IF(MATCH(CD{r},$CD$4:$CD${BUY_END},0)=ROW()-3,1,0))')
        put(f'CF{r}', f'=N(CF{r - 1})+CE{r}')
    put('CA1', f'=MAX($CB$4:$CB${FIN_END})')
    put('CD1', f'=MAX($CF$4:$CF${BUY_END})')
    put('CJ1', f'=MAX($CK$4:$CK${FIN_END})')
    put('CM1', f'=MAX($CN$4:$CN${FIN_END})')
    compact(ws, 'CG', ['CH'], ['CD'], 'CF', '$CD$1', BUY_END, 123)

    # CP～CW：原料出库明细逐行（公司购买命中、估算出库重量）
    for r in range(4, OUT_END + 1):
        F_, D_, B2, N_, H2, U_, W_ = (IX(OUT_S, c, OUT_END, p) for c in 'FDBNHUW')
        put(f'CV{r}', f'=IF({F_}="","",{F_})')
        put(f'CW{r}', f'=IF($CV{r}="","",{D_})')
        put(f'CU{r}', f'=IF($CV{r}="","",{date_num(B2)})')
        put(f'CP{r}', f'=IF(AND($CV{r}<>"",{N_}="公司购买"),1,0)')
        put(f'CQ{r}', f'=N(CQ{r - 1})+CP{r}')
        put(f'CR{r}', f'=IF(CP{r}=0,0,IF(AND(N($CU{r})>=果然鲜采购明细!$Q$1,N($CU{r})<=果然鲜采购明细!$Q$2),1,0))')
        put(f'CS{r}', f'=N(CS{r - 1})+CR{r}')
        # 公司购买金额＝出库数量×采购单价，在这里现算（不靠原料出库明细 V 列：中间插的行没有 V 公式也照样算对）
        put(f'CX{r}', f'=IF(CP{r}=1,ROUND(N({H2})*N({U_}),2),0)')
        put(f'CY{r}', f'=IF(CP{r}=1,{W_}&"","")')
        put(f'CT{r}', f'=IF($CV{r}="","",N({H2})*N(IFERROR(INDEX(库存结余!$I$4:$I${KC_END},'
                      f'MATCH($A{r + P_OUT0 - 4},$K$4:$K${KC_END},0)),0)))')
    put('CP1', f'=MAX($CQ$4:$CQ${OUT_END})')
    put('CR1', f'=MAX($CS$4:$CS${OUT_END})')


def compact(ws, poscol, outcols, srccols, cumcol, total, pool_end, last_row):
    """紧凑清单：位置列＝MATCH(第 n 个, 累计列)，超过个数就不查；取数列按位置 INDEX"""
    for r in range(4, last_row + 1):
        ws[f'{poscol}{r}'] = f'=IF(ROW()-3>{total},"",MATCH(ROW()-3,${cumcol}$4:${cumcol}${pool_end},0))'
        for c, s in zip(outcols, srccols):
            ws[f'{c}{r}'] = f'=IF(${poscol}{r}="","",INDEX(${s}$4:${s}${pool_end},${poscol}{r}))'


# ════════════════════════════════════════════════════════════════════════════
# ③ 库存结余：插 I（平均重量）、P/Q/R（在库成品、损耗、损耗率）
# ════════════════════════════════════════════════════════════════════════════
def _move_cols(ws, mapping, max_row, max_col):
    """按 旧列号→新列号 搬单元格（值＋样式），从右往左搬，空出来的格子清掉"""
    cells = {}
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            x = ws.cell(row=r, column=c)
            cells[(r, c)] = (x.value, copy.copy(x._style) if x.has_style else None)
    for m in [str(m) for m in ws.merged_cells.ranges]:
        ws.unmerge_cells(m)
    for (r, c) in cells:
        x = ws.cell(row=r, column=c)
        x.value = None
    for (r, c), (v, st) in cells.items():
        nc = mapping(c)
        x = ws.cell(row=r, column=nc)
        x.value = v
        if st is not None:
            x._style = st
    dims = []
    for k, d in ws.column_dimensions.items():
        lo, hi = (d.min or CI(k)), (d.max or CI(k))
        for c in range(lo, hi + 1):
            dims.append((c, d.width, d.hidden))
    for k in list(ws.column_dimensions):
        del ws.column_dimensions[k]
    for c, w, h in dims:
        nk = L(mapping(c))
        if w:
            ws.column_dimensions[nk].width = w
        if h:
            ws.column_dimensions[nk].hidden = h


def stock(wb):
    ws = wb['库存结余']
    T = 404
    mp = lambda c: c if c <= 8 else (c + 1 if c <= 14 else c + 4)
    _move_cols(ws, mp, T, 21)            # A～U → A～H、J～O、S～Y
    # 第 2 行：说明合并 A2:P2，日期块挪到 Q2:T2，隐藏辅助 X1/X2
    lab1, in1, lab2, in2 = (copy.copy(ws[a]._style) for a in ('N2', 'O2', 'S2', 'T2'))
    v1, v2 = ws['O2'].value, ws['T2'].value
    for a in ('N2', 'O2', 'S2', 'T2'):
        ws[a].value = None
    for a in ('N2', 'O2'):
        ws[a]._style = copy.copy(ws['M2']._style)
    ws['Q2'], ws['R2'], ws['S2'], ws['T2'] = '开始日期', v1, '结束日期', v2
    ws['Q2']._style, ws['R2']._style, ws['S2']._style, ws['T2']._style = lab1, in1, lab2, in2
    ws['X1'] = '=IF($R2="",DATE(1900,1,1),$R2)'
    ws['X2'] = '=IF($T2="",DATE(2999,12,31),$T2)'
    ws.merge_cells('A1:U1')
    ws.merge_cells('A2:P2')
    ws['A2'] = ('★ 全自动。填右边【开始日期／结束日期】按区间统计（留空＝不限）。平均重量＝全部入库总重÷入库数量（不随日期）；'
                '出库总重＝出库数量×平均重量；在库成品（含损耗）＝加工类出库（出库加工、二次加工）数量×平均重量－成品出库总重；'
                '损耗只在原料出完（结余数量＝0）后才算＝入库总重－直接出库数量×平均重量－成品出库总重（报损算损耗）。'
                '负数标红：多半是加工时重新分级，按货主＋品种合起来看。')
    ws.row_dimensions[2].height = 45

    IN = lambda col: R(IN_S, col, IN_END)
    OUT = lambda col: R(OUT_S, col, OUT_END)
    FIN = lambda col: R(FIN_S, col, FIN_END)

    def k3(rng, r):
        return f'{rng("F")},$B{r},{rng("D")},$C{r},{rng("E")},$D{r}'

    def dt(rng):
        return f'{rng("B")},">="&$X$1,{rng("B")},"<="&$X$2'

    hdr = {'I': '平均重量\nKG/件', 'P': '在库成品吨位\n（含损耗）KG', 'Q': '损耗KG', 'R': '损耗率',
           'Y': '加工出库数量\n（辅助）', 'Z': '报损数量\n（辅助）'}
    for c, t in hdr.items():
        ws[f'{c}3'] = t
        style_from(ws[f'{c}3'], ws['H3'])
    for r in range(4, T):
        ws[f'G{r}'] = f'=IF($B{r}="","",SUMIFS({IN("H")},{k3(IN, r)},{dt(IN)}))'
        ws[f'H{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({IN("K")},{k3(IN, r)},{dt(IN)}),2))'
        ws[f'I{r}'] = (f'=IF($B{r}="","",IFERROR(SUMIFS({IN("K")},{k3(IN, r)})/SUMIFS({IN("H")},{k3(IN, r)}),""))')
        ws[f'J{r}'] = f'=IF($B{r}="","",SUMIFS({OUT("H")},{k3(OUT, r)},{dt(OUT)}))'
        ws[f'K{r}'] = f'=IF($B{r}="","",ROUND(N($J{r})*N($I{r}),2))'
        ws[f'L{r}'] = f'=IF($B{r}="","",N($G{r})-N($J{r}))'
        ws[f'M{r}'] = f'=IF($B{r}="","",ROUND(N($H{r})-N($K{r}),2))'
        ws[f'N{r}'] = f'=IF($B{r}="","",SUMIFS({FIN("J")},{k3(FIN, r)},{dt(FIN)}))'
        ws[f'O{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({FIN("M")},{k3(FIN, r)},{dt(FIN)}),2))'
        ws[f'P{r}'] = f'=IF($B{r}="","",ROUND(N($Y{r})*N($I{r})-N($O{r}),2))'
        ws[f'Q{r}'] = (f'=IF($B{r}="","",IF(N($L{r})<>0,"",'
                       f'ROUND(N($H{r})-(N($J{r})-N($Y{r})-N($Z{r}))*N($I{r})-N($O{r}),2)))')
        ws[f'R{r}'] = f'=IF($B{r}="","",IF(OR($Q{r}="",N($Y{r})*N($I{r})=0),"",ROUND(N($Q{r})/(N($Y{r})*N($I{r})),4)))'
        ws[f'S{r}'] = (f'=IF($B{r}="","",IF($L{r}<=0,"无库存",IF($L{r}>=基础资料!$V$4,"充足",'
                       f'IF($L{r}>=基础资料!$V$5,"正常","偏低"))))')
        ws[f'T{r}'] = (f'=IF($B{r}="","",IF(N($L{r})<>0,"原料未出完","")&IF(AND(N($L{r})<>0,N($P{r})<0),"；","")'
                       f'&IF(N($P{r})<0,"本等级成品多于加工出库（可能加工时重新分级）",""))')
        ws[f'Y{r}'] = (f'=IF($B{r}="","",SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"出库加工",{dt(OUT)})'
                       f'+SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"二次加工",{dt(OUT)}))')
        ws[f'Z{r}'] = f'=IF($B{r}="","",SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"报损",{dt(OUT)}))'
        # 样式：新列照 H（重量）／G（数量）
        for c, src, fmt in (('I', 'H', '0.000'), ('P', 'H', MONEY_FMT), ('Q', 'H', MONEY_FMT), ('R', 'H', PCT_FMT),
                            ('K', 'H', MONEY_FMT), ('M', 'H', MONEY_FMT)):
            style_from(ws[f'{c}{r}'], ws[f'{src}{r}'])
            ws[f'{c}{r}'].number_format = fmt
        style_from(ws[f'T{r}'], ws[f'S{r}'])
        ws[f'T{r}'].number_format = 'General'
        ws[f'T{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
        for c in 'YZ':
            ws[f'{c}{r}'].font = F_HELP
    # 合计行
    for c in 'GHJKLMNOPQ':
        ws[f'{c}{T}'] = f'=ROUND(SUM({c}4:{c}{T - 1}),2)'
        style_from(ws[f'{c}{T}'], ws[f'H{T}'])
        ws[f'{c}{T}'].number_format = '#,##0;[Red]\\-#,##0;\\-' if c in 'GJLN' else MONEY_FMT
    for c in 'IRST':
        ws[f'{c}{T}'].value = None
        style_from(ws[f'{c}{T}'], ws[f'B{T}'])
    for c in 'YZ':
        ws[f'{c}3'].font = F_HELP
    for c, w in {'I': 11, 'P': 15, 'Q': 11, 'R': 9, 'T': 30, 'Y': 10, 'Z': 10}.items():
        ws.column_dimensions[c].width = w
    for c in ('Y', 'Z'):
        ws.column_dimensions[c].hidden = True
    ws.auto_filter.ref = f'A3:T{T}'
    ws.print_area = 'A1:U20'


def stock_total(wb):
    """库存总结余 H（出库总重）改成按等级平均重量估：SUMIFS(_自动清单 估算出库重量)，跟库存结余同口径、照样按本表日期筛"""
    ws = wb['库存总结余']
    A = '_自动清单!'
    for r in range(4, 304):
        ws[f'H{r}'] = (f'=IF($B{r}="","",ROUND(SUMIFS({A}$CT$4:$CT${OUT_END},{A}$CV$4:$CV${OUT_END},$B{r},'
                       f'{A}$CW$4:$CW${OUT_END},$C{r},{A}$CU$4:$CU${OUT_END},">="&$P$1,{A}$CU$4:$CU${OUT_END},"<="&$P$2),2))')
    ws['H3'] = '出库总重KG\n（数量×平均重量）'
    a2 = ws['A2'].value or ''
    if '平均重量' not in a2:
        ws['A2'] = a2 + ' 出库总重＝每笔出库数量×该等级平均重量（跟【库存结余】同口径）。'


def stock_check(wb):
    """实盘库存盘点 H（账面结存总重）＝入库总重－出库数量×平均重量（跟库存结余同口径）"""
    ws = wb['实盘库存盘点']
    for r in range(5, 305):
        k3i = (f'{R(IN_S, "F", IN_END)},$B{r},{R(IN_S, "D", IN_END)},$C{r},{R(IN_S, "E", IN_END)},$D{r},'
               f'{R(IN_S, "B", IN_END)},">="&$Q$1,{R(IN_S, "B", IN_END)},"<="&$Q$2')
        k3o = (f'{R(OUT_S, "F", OUT_END)},$B{r},{R(OUT_S, "D", OUT_END)},$C{r},{R(OUT_S, "E", OUT_END)},$D{r},'
               f'{R(OUT_S, "B", OUT_END)},">="&$Q$1,{R(OUT_S, "B", OUT_END)},"<="&$Q$2')
        ws[f'H{r}'] = (f'=IF($B{r}="","",ROUND(SUMIFS({R(IN_S, "K", IN_END)},{k3i})-SUMIFS({R(OUT_S, "H", OUT_END)},{k3o})'
                       f'*N(INDEX(库存结余!$I$4:$I${KC_END},ROW()-4)),2))')


# ════════════════════════════════════════════════════════════════════════════
# ④ 库存等级接口（K8）：不随库存结余的日期筛选
# ════════════════════════════════════════════════════════════════════════════
def stock_iface(wb):
    ws = wb['库存等级接口']
    IN = lambda col: R(IN_S, col, IN_END)
    OUT = lambda col: R(OUT_S, col, OUT_END)
    FIN = lambda col: R(FIN_S, col, FIN_END)

    def k3(rng, r):
        return f'{rng("F")},$B{r},{rng("D")},$C{r},{rng("E")},$D{r}'

    st_lab, st_num = copy.copy(ws['J1']._style), copy.copy(ws['K1']._style)
    ws['J1'].value = None
    ws['K1'].value = None
    ws['R1'], ws['S1'] = '有效行数', f'=COUNTIF($B$4:$B${KC_END},"?*")'
    ws['R1']._style, ws['S1']._style = st_lab, st_num
    for c, t in {'I': '平均重量\nKG/件', 'J': '在库成品吨位\n（含损耗）KG', 'K': '成品出库总重KG',
                 'L': '损耗KG', 'M': '加工出库重量KG'}.items():
        ws[f'{c}3'] = t
        style_from(ws[f'{c}3'], ws['H3'])
    ws.row_dimensions[3].height = 32
    for r in range(4, KC_END + 1):
        p = 'ROW()-3'
        for c, src in (('B', 'B'), ('C', 'C'), ('D', 'D'), ('E', 'F')):
            x = f'INDEX(库存结余!${src}$4:${src}${KC_END},{p})'
            ws[f'{c}{r}'] = f'=IF({x}="","",{x})'
        ws[f'F{r}'] = f'=IF($B{r}="","",SUMIFS({IN("H")},{k3(IN, r)})-SUMIFS({OUT("H")},{k3(OUT, r)}))'
        ws[f'G{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({IN("K")},{k3(IN, r)})-SUMIFS({OUT("H")},{k3(OUT, r)})*N($I{r}),2))'
        ws[f'H{r}'] = (f'=IF($B{r}="","",IF(N($F{r})<=0,"无库存",IF($F{r}>=基础资料!$V$4,"充足",'
                       f'IF($F{r}>=基础资料!$V$5,"正常","偏低"))))')
        x = f'INDEX(库存结余!$I$4:$I${KC_END},{p})'
        ws[f'I{r}'] = f'=IF($B{r}="","",IF(N({x})=0,"",{x}))'
        ws[f'M{r}'] = (f'=IF($B{r}="","",ROUND((SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"出库加工")'
                       f'+SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"二次加工"))*N($I{r}),2))')
        ws[f'K{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({FIN("M")},{k3(FIN, r)}),2))'
        ws[f'J{r}'] = f'=IF($B{r}="","",ROUND(N($M{r})-N($K{r}),2))'
        ws[f'L{r}'] = (f'=IF($B{r}="","",IF(N($F{r})<>0,"",ROUND(N($G{r})+N($J{r})'
                       f'+SUMIFS({OUT("H")},{k3(OUT, r)},{OUT("N")},"报损")*N($I{r}),2)))')
        for c in 'IJKLM':
            style_from(ws[f'{c}{r}'], ws[f'G{r}'])
            ws[f'{c}{r}'].number_format = '0.000' if c == 'I' else MONEY_FMT
    for m in [str(m) for m in ws.merged_cells.ranges]:
        if m in ('A1:H1', 'A2:H2'):
            ws.unmerge_cells(m)
            ws.merge_cells(m.replace('H', 'M'))
    ws['A2'] = ('★ 全自动，1:1 对应【库存结余】的行（货主＋品种＋等级），但按全部日期算，不跟【库存结余】的日期筛选走。'
                '《03》按固定位置引用本表 A3:M403，请勿插入/删除行列、改表名。I 平均重量＝入库总重÷入库数量；'
                'J 在库成品（含损耗）＝M 加工出库重量－K 成品出库总重；L 损耗：原料没出完留空，出完后＝在库成品＋报损重量。')
    ws.row_dimensions[2].height = 45
    for c, w in {'I': 11, 'J': 14, 'K': 13, 'L': 11, 'M': 13, 'R': 9, 'S': 7}.items():
        ws.column_dimensions[c].width = w


# ════════════════════════════════════════════════════════════════════════════
# ⑦ 对账明细接口（K1）：6,000 行，原料出库「公司购买」也出采购字段
# ════════════════════════════════════════════════════════════════════════════
def recon(wb):
    ws = wb['对账明细接口']
    SEG = {'IN': (IN_S, IN_END, 0), 'OUT': (OUT_S, OUT_END, T_IN), 'FIN': (FIN_S, FIN_END, T_OUT),
           'DAI': (DAI_S, DAI_END, T_FIN)}

    def pz(seg, r):
        off = SEG[seg][2]
        return f'$Z{r}' if off == 0 else f'$Z{r}-{off}'

    def v(seg, col, r):
        sh, end, _ = SEG[seg]
        return V(sh, col, end, pz(seg, r))

    def nv(seg, col, r):
        sh, end, _ = SEG[seg]
        return NV(sh, col, end, pz(seg, r))

    def ix(seg, col, r):
        sh, end, _ = SEG[seg]
        return IX(sh, col, end, pz(seg, r))

    def chain(r, e_in, e_out, e_fin, e_dai):
        e = [x if x is not None else '""' for x in (e_in, e_out, e_fin, e_dai)]
        return (f'=IF($Z{r}="","",IF($Z{r}<={T_IN},{e[0]},IF($Z{r}<={T_OUT},{e[1]},'
                f'IF($Z{r}<={T_FIN},{e[2]},{e[3]}))))')

    def buy(r, e):
        """原料出库：只有出库原因＝公司购买 的行才出采购字段"""
        return f'IF({ix("OUT", "N", r)}<>"公司购买","",{e})'

    def fin_out(r, e_out, e_fin):
        return f'=IF(OR($Z{r}="",$Z{r}<={T_IN},$Z{r}>{T_FIN}),"",IF($Z{r}<={T_OUT},{e_out},{e_fin}))'

    def fin_only(r, e_fin):
        return f'=IF(OR($Z{r}="",$Z{r}<={T_OUT},$Z{r}>{T_FIN}),"",{e_fin})'

    def gm(r):
        x = f'N(INDEX(_自动清单!$CX$4:$CX${OUT_END},{pz("OUT", r)}))'
        return f'IF({x}=0,"",{x})'

    tmpl = {L(c): ws.cell(row=1203, column=c) for c in range(1, 33)}
    ct = '_自动清单!$CT$4:$CT$' + str(OUT_END)
    for r in range(4, DZ_END + 1):
        if r > 1203:
            for c, s in tmpl.items():
                if s.has_style:
                    ws[f'{c}{r}']._style = copy.copy(s._style)
        f = {}
        f['A'] = f'=IF($B{r}="","",ROW()-3)'
        f['B'] = chain(r, v('IN', 'B', r), v('OUT', 'B', r), v('FIN', 'B', r), v('DAI', 'B', r))
        f['C'] = chain(r, v('IN', 'F', r), v('OUT', 'F', r), v('FIN', 'F', r), v('DAI', 'G', r))
        f['D'] = chain(r, '"原料入库"', f'"原料出库－"&{ix("OUT", "N", r)}', f'"成品出库－"&{ix("FIN", "I", r)}',
                       f'"其他代存－"&{ix("DAI", "D", r)}')
        f['E'] = chain(r, *(v(s, 'C', r) for s in ('IN', 'OUT', 'FIN', 'DAI')))
        f['F'] = chain(r, v('IN', 'D', r), v('OUT', 'D', r), v('FIN', 'D', r), v('DAI', 'E', r))
        f['G'] = chain(r, v('IN', 'E', r), v('OUT', 'E', r), v('FIN', 'E', r), v('DAI', 'F', r))
        f['H'] = chain(r, v('IN', 'H', r), v('OUT', 'H', r), v('FIN', 'J', r), v('DAI', 'H', r))
        f['I'] = chain(r, v('IN', 'I', r), v('OUT', 'I', r), v('FIN', 'K', r), v('DAI', 'I', r))
        est = f'N(INDEX({ct},{pz("OUT", r)}))'
        j_out = (f'IF(N({ix("OUT", "K", r)})<>0,N({ix("OUT", "K", r)}),IF({est}=0,"",ROUND({est},2)))')
        f['J'] = chain(r, v('IN', 'K', r), j_out, v('FIN', 'M', r), v('DAI', 'K', r))
        f['K'] = fin_only(r, v('FIN', 'AI', r))
        f['L'] = fin_only(r, v('FIN', 'AD', r))
        f['M'] = chain(r, v('IN', 'Q', r), v('OUT', 'T', r), v('FIN', 'AF', r), v('DAI', 'N', r))
        f['N'] = chain(r, v('IN', 'L', r), v('OUT', 'L', r), v('FIN', 'O', r), None)
        f['O'] = chain(r, None, nv('OUT', 'M', r), nv('FIN', 'P', r), None)
        f['P'] = '=""'
        f['Q'] = f'=IF(OR($Z{r}="",$Z{r}>{T_IN}),"",{nv("IN", "M", r)})'
        f['R'] = chain(r, v('IN', 'N', r), v('OUT', 'Q', r), v('FIN', 'Q', r), None)
        f['S'] = f'=IF($N{r}="","",IF(LEFT($N{r},2)="自备","自备筐","公司筐"))'
        f['T'] = fin_out(r, buy(r, v('OUT', 'F', r)), v('FIN', 'F', r))
        f['U'] = fin_out(r, buy(r, nv('OUT', 'U', r)), v('FIN', 'T', r))
        f['V'] = fin_out(r, buy(r, gm(r)), v('FIN', 'V', r))
        f['W'] = fin_out(r, buy(r, v('OUT', 'W', r)), v('FIN', 'X', r))
        f['X'] = fin_out(r, buy(r, v('OUT', 'X', r)),
                         f'IF({ix("FIN", "H", r)}<>"",{ix("FIN", "H", r)},{v("FIN", "G", r)})')
        f['Y'] = fin_only(r, nv('FIN', 'Y', r))
        f['Z'] = f'=IF(ROW()-3>_自动清单!$BX$1,"",MATCH(ROW()-3,_自动清单!$BY$4:$BY${P_END4},0))'
        f['AA'] = fin_only(r, nv('FIN', 'Z', r))
        f['AB'] = fin_only(r, nv('FIN', 'AA', r))
        f['AC'] = fin_out(r, '""', nv('FIN', 'U', r))
        f['AD'] = fin_only(r, v('FIN', 'I', r))
        f['AE'] = fin_only(r, nv('FIN', 'AB', r))
        f['AF'] = fin_out(r, buy(r, gm(r)), nv('FIN', 'AJ', r))
        for c, x in f.items():
            ws[f'{c}{r}'] = x
    ws['A2'] = (f'=IF(N(_自动清单!$BX$1)>{DZ_END - 3},"★★ 注意：业务笔数已有 "&N(_自动清单!$BX$1)&" 个，超过本表 '
                f'{DZ_END - 3} 个的显示上限，请联系维护人员扩容（《04》也要一起扩）","★《04》按固定位置引用本表 A3:AF{DZ_END}，'
                f'请勿插入/删除行列或改表名。原料出库「公司购买」的行也带采购单价/金额/付款状态（T～X、AF）。'
                f'【业务笔数 "&N(_自动清单!$BX$1)&"/{DZ_END - 3}】")')
    # 条件格式：超过上限变色
    cf = ws.conditional_formatting
    for rng in list(cf._cf_rules):
        for rule in cf._cf_rules[rng]:
            if rule.formula:
                rule.formula = [x.replace('>1200', f'>{DZ_END - 3}') for x in rule.formula]


# ════════════════════════════════════════════════════════════════════════════
# ⑥ 果然鲜采购明细：成品出库来源 ＋ 原料出库「公司购买」来源
# ════════════════════════════════════════════════════════════════════════════
def purchase(wb):
    ws = wb['果然鲜采购明细']
    A = '_自动清单!'
    # P＝来源（显示），R＝源行（隐藏）；Q1/Q2 日期辅助照旧
    ws['P3'] = '来源'
    style_from(ws['P3'], ws['O3'])
    ws['R3'] = '源位置'
    ws['R3'].font = F_HELP
    ws.column_dimensions['P'].hidden = False
    ws.column_dimensions['P'].width = 10
    ws.column_dimensions['R'].hidden = True
    F = lambda col, r: IX(FIN_S, col, FIN_END, f'$R{r}')
    O = lambda col, r: IX(OUT_S, col, OUT_END, f'$R{r}')
    fv = lambda col, r: V(FIN_S, col, FIN_END, f'$R{r}')
    ov = lambda col, r: V(OUT_S, col, OUT_END, f'$R{r}')
    onv = lambda col, r: NV(OUT_S, col, OUT_END, f'$R{r}')
    ogm = lambda r: NV('_自动清单', 'CX', OUT_END, f'$R{r}')
    for r in range(4, 404):
        n = f'ROW()-3'
        ws[f'R{r}'] = (f'=IF({n}<={A}$CJ$1,MATCH({n},{A}$CK$4:$CK${FIN_END},0),'
                       f'IF({n}-{A}$CJ$1<={A}$CR$1,MATCH({n}-{A}$CJ$1,{A}$CS$4:$CS${OUT_END},0),""))')
        ws[f'P{r}'] = f'=IF($R{r}="","",IF(ROW()-3<={A}$CJ$1,"成品出库","原料出库"))'
        style_from(ws[f'P{r}'], ws[f'H{r}'])
        ws[f'R{r}'].font = F_HELP
        sel = lambda e_out, e_fin: f'=IF($R{r}="","",IF($P{r}="原料出库",{e_out},{e_fin}))'
        ws[f'A{r}'] = f'=IF($R{r}="","",ROW()-3)'
        ws[f'B{r}'] = sel(ov('B', r), fv('B', r))
        ws[f'C{r}'] = sel(ov('C', r), fv('C', r))
        ws[f'D{r}'] = sel(ov('F', r), fv('F', r))
        ws[f'E{r}'] = sel(ov('D', r), fv('D', r))
        ws[f'F{r}'] = sel(ov('E', r), fv('E', r))
        ws[f'G{r}'] = sel(ov('H', r), fv('J', r))
        ws[f'H{r}'] = sel(ov('I', r), fv('K', r))
        ws[f'I{r}'] = sel(onv('U', r), fv('T', r))
        ws[f'J{r}'] = sel('""', fv('U', r))
        ws[f'K{r}'] = sel(ogm(r), fv('V', r))
        ws[f'L{r}'] = sel(ov('W', r), fv('X', r))
        m_out = (f'{ov("T", r)}&IF(N({O("U", r)})=0,IF({O("T", r)}=0,"","；")&"未填采购单价","")')
        m_fin = (f'{fv("AF", r)}&IF(AND({F("F", r)}="果然鲜",N({F("T", r)})<>0),IF({F("AF", r)}=0,"","；")'
                 f'&"货主是果然鲜自己，确认是不是该填在原料出库","")')
        ws[f'M{r}'] = sel(m_out, m_fin)
        ws[f'N{r}'] = sel(ov('X', r), f'IF({F("H", r)}<>"",{F("H", r)},{fv("G", r)})')
        ws[f'O{r}'] = sel(ogm(r), f'IF(N({F("AJ", r)})=0,"",N({F("AJ", r)}))')
    # 底部按货主汇总：523～527 行跟上面统一
    for r in range(408, 528):
        ws[f'A{r}'] = f'=IF($B{r}="","",ROW()-407)'
        ws[f'B{r}'] = f'=IFERROR(INDEX(_自动清单!$BO$4:$BO$1203,ROW()-407),"")'
        ws[f'C{r}'] = f'=IF($B{r}="","",COUNTIF($D$4:$D$403,$B{r}))'
        ws[f'D{r}'] = f'=IF($B{r}="","",SUMIF($D$4:$D$403,$B{r},$G$4:$G$403))'
        ws[f'E{r}'] = f'=IF($B{r}="","",ROUND(SUMIF($D$4:$D$403,$B{r},$J$4:$J$403),2))'
        ws[f'F{r}'] = f'=IF($B{r}="","",ROUND(SUMIF($D$4:$D$403,$B{r},$K$4:$K$403),2))'
        ws[f'G{r}'] = f'=IF($B{r}="","",ROUND(N($E{r})+N($F{r}),2))'
        ws[f'H{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS($O$4:$O$403,$D$4:$D$403,$B{r},$L$4:$L$403,"已付清"),2))'
        ws[f'I{r}'] = f'=IF($B{r}="","",ROUND(N($G{r})-N($H{r}),2))'
        for c in 'GHI':
            if not ws[f'{c}{r}'].has_style or r >= 523:
                style_from(ws[f'{c}{r}'], ws[f'{c}408'])
    ws['I528'] = '=ROUND(SUM(I408:I527),2)'
    style_from(ws['I528'], ws['H528'])
    ws['A2'] = ('按【成品出库明细】填了采购单价/采购金额的行，加上【原料出库明细】出库原因＝「公司购买」的行（P 列标来源），'
                '没填单价的也列出来、备注写「未填采购单价」。采购金额只是货款，应付合计另加采购代发运费。右边日期筛选。')
    ws.row_dimensions[2].height = 36
    ws.auto_filter.ref = 'A3:P404'


# ════════════════════════════════════════════════════════════════════════════
# ⑩ 果然鲜销售明细：去 INDIRECT、修二段、接「三、按购买方汇总」；删果然鲜销售汇总
# ════════════════════════════════════════════════════════════════════════════
S3_T, S3_H, S3_0, S3_1, S3_SUM, S3_NOTE = 478, 479, 480, 599, 600, 602


def sales(wb):
    ws = wb['果然鲜销售明细']
    old = wb['果然鲜销售汇总']
    A = '_自动清单!'
    cols = {'B': 'B', 'C': 'C', 'D': 'I', 'E': 'F', 'G': 'D', 'H': 'E', 'I': 'J', 'J': 'K', 'K': 'M', 'L': 'T',
            'M': 'V', 'N': 'Y', 'O': 'AB', 'P': 'Z', 'Q': 'AA', 'R': 'AC', 'S': 'AD', 'T': 'AF'}
    for r in range(4, 404):
        ws[f'W{r}'] = f'=IF(ROW()-3>{A}$CA$1,"",MATCH(ROW()-3,{A}$CB$4:$CB${FIN_END},0))'
        for c, src in cols.items():
            ws[f'{c}{r}'] = f'=IF($W{r}="","",{V(FIN_S, src, FIN_END, f"$W{r}")})'
        h = IX(FIN_S, 'H', FIN_END, f'$W{r}')
        ws[f'F{r}'] = f'=IF($W{r}="","",IF({h}<>"",{h},{V(FIN_S, "G", FIN_END, f"$W{r}")}))'
        ws[f'U{r}'] = f'=IF($W{r}="","",{NV(FIN_S, "AI", FIN_END, f"$W{r}")})'
        # 二段辅助：Z 货主首次、AA 累计、AB 货主＋品种首次（算品种数）
        ws[f'Z{r}'] = f'=IF($E{r}="",0,IF(MATCH($E{r},$E$4:$E$403,0)=ROW()-3,1,0))'
        ws[f'AA{r}'] = f'=N(AA{r - 1})+Z{r}'
        ws[f'AB{r}'] = f'=IF($E{r}="",0,IF(COUNTIFS($E$4:$E{r},$E{r},$G$4:$G{r},$G{r})=1,1,0))'
        for c in ('Z', 'AA', 'AB'):
            ws[f'{c}{r}'].font = F_HELP
    ws['Z3'], ws['AA3'], ws['AB3'] = '货主首次', '累计', '货主品种首次'
    ws['AA1'] = '=MAX($AA$4:$AA$403)'
    for c in ('Z3', 'AA3', 'AB3', 'AA1'):
        ws[c].font = F_HELP
    for c in ('Z', 'AA', 'AB'):
        ws.column_dimensions[c].hidden = True
    # 二、按货主/客户（416～475）：扫全部 400 行明细；品种数改真的品种数
    for r in range(416, 476):
        for c in ('Z', 'AA', 'AB'):
            ws[f'{c}{r}'].value = None
        n = f'ROW()-415'
        ws[f'B{r}'] = f'=IF({n}>$AA$1,"",INDEX($E$4:$E$403,MATCH({n},$AA$4:$AA$403,0)))'
        ws[f'C{r}'] = f'=IF($B{r}="","",SUMIFS($AB$4:$AB$403,$E$4:$E$403,$B{r}))'
    ws.column_dimensions['J'].width = 13

    # 三、按购买方汇总（搬自果然鲜销售汇总）
    ws.merge_cells(f'A{S3_T}:T{S3_T}')
    ws[f'A{S3_T}'] = '三、按 购 买 方 汇 总（卖给谁、收了多少；对应的采购＝成品出库采购＋原料出库「公司购买」按购买方/去向）'
    for c in range(1, 21):
        style_from(ws.cell(row=S3_T, column=c), ws.cell(row=414, column=c))
    ws.row_dimensions[S3_T].height = ws.row_dimensions[414].height
    for c in range(1, 18):
        L_ = L(c)
        ws[f'{L_}{S3_H}'] = old[f'{L_}3'].value
        style_from(ws[f'{L_}{S3_H}'], ws[f'{"ABCDEFGHI"[min(c, 9) - 1]}415'])
    ws.row_dimensions[S3_H].height = 30
    sale = lambda col: f'${col}$4:${col}$403'
    P = '果然鲜采购明细!'
    for r in range(S3_0, S3_1 + 1):
        n = f'ROW()-{S3_H}'
        f = {
            'A': f'=IF($B{r}="","",{n})',
            'B': f'=IF({n}>{A}$CD$1,"",INDEX({A}$CH$4:$CH$123,{n}))',
            'C': f'=IF($B{r}="","",COUNTIF({sale("F")},$B{r}))',
            'D': f'=IF($B{r}="","",SUMIF({sale("F")},$B{r},{sale("I")}))',
            'E': f'=IF($B{r}="","",ROUND(SUMIF({sale("F")},$B{r},{sale("K")}),2))',
            'F': f'=IF($B{r}="","",ROUND(SUMIF({sale("F")},$B{r},{sale("O")}),2))',
            'G': f'=IF($B{r}="","",ROUND(SUMIF({sale("F")},$B{r},{sale("P")}),2))',
            'H': f'=IF($B{r}="","",ROUND(SUMIF({sale("F")},$B{r},{sale("Q")}),2))',
            'I': f'=IF($B{r}="","",ROUND(N($F{r})+N($G{r})+N($H{r}),2))',
            'J': f'=IF($B{r}="","",ROUND(SUMIFS({sale("U")},{sale("F")},$B{r},{sale("S")},"已收讫"),2))',
            'K': f'=IF($B{r}="","",ROUND(N($I{r})-N($J{r}),2))',
            'L': f'=IF($B{r}="","",SUMIFS({P}$G$4:$G$403,{P}$N$4:$N$403,$B{r}))',
            'M': f'=IF($B{r}="","",SUMIFS({P}$K$4:$K$403,{P}$N$4:$N$403,$B{r}))',
            'N': f'=IF($B{r}="","",SUMIFS({P}$J$4:$J$403,{P}$N$4:$N$403,$B{r}))',
            'O': f'=IF($B{r}="","",SUMIFS({P}$O$4:$O$403,{P}$N$4:$N$403,$B{r}))',
            'P': f'=IF($B{r}="","",SUMIFS({P}$O$4:$O$403,{P}$N$4:$N$403,$B{r},{P}$L$4:$L$403,"已付清"))',
            'Q': f'=IF($B{r}="","",ROUND(N($O{r})-N($P{r}),2))',
        }
        for c, x in f.items():
            ws[f'{c}{r}'] = x
            style_from(ws[f'{c}{r}'], old[f'{c}4'])
    ws[f'A{S3_SUM}'] = '合  计'
    for c in range(1, 18):
        L_ = L(c)
        style_from(ws[f'{L_}{S3_SUM}'], old[f'{L_}124'])
        if c >= 3:
            ws[f'{L_}{S3_SUM}'] = f'=ROUND(SUM({L_}{S3_0}:{L_}{S3_1}),2)'
    ws.merge_cells(f'A{S3_NOTE}:Q{S3_NOTE}')
    ws[f'A{S3_NOTE}'] = (old['A126'].value + '「采购数量／采购货款／应付合计」取【果然鲜采购明细】里「购买方(再卖给谁)」是这家的行'
                         '（含原料出库「公司购买」按「购买方/去向」），跟着采购明细的日期筛选；筐、箱数量直接相加，只作参考。')
    style_from(ws[f'A{S3_NOTE}'], old['A126'])
    ws.row_dimensions[S3_NOTE].height = 45
    a2 = ws['A2'].value or ''
    if '第 478 行' not in a2:
        ws['A2'] = a2 + ('　★ 往下翻：第 406 行「一、按销售类型」，第 414 行「二、按货主/客户」，'
                         '第 478 行「三、按购买方」（原【果然鲜销售汇总】已并到这里）。')
    ws.row_dimensions[2].height = 45
    for c, w in {'B': 14, 'L': 11, 'Q': 11}.items():
        if (ws.column_dimensions[c].width or 0) < w:
            ws.column_dimensions[c].width = w
    wb.remove(old)


# ════════════════════════════════════════════════════════════════════════════
# ⑧ 财务取数接口第三块：公司购买果品款加上原料出库「公司购买」
# ════════════════════════════════════════════════════════════════════════════
def finance(wb):
    ws = wb['财务取数接口']
    o = lambda col: R('_自动清单', col, OUT_END)
    f = lambda col: R(FIN_S, col, FIN_END)
    for r in range(207, 507):
        ws[f'D{r}'] = (f'=IF(OR($B{r}="",$C{r}<>"成品销售款"),0,ROUND(SUMIFS({f("AI")},{f("AK")},$B{r}),2))')
        ws[f'E{r}'] = (f'=IF(OR($B{r}="",$C{r}<>"公司购买果品款"),0,ROUND(SUMIFS({f("AJ")},{f("F")},$B{r})'
                       f'+SUMIFS({o("CX")},{o("CV")},$B{r}),2))')
        ws[f'H{r}'] = (f'=IF($B{r}="","",IF($C{r}="成品销售款",ROUND(SUMIFS({f("AI")},{f("AK")},$B{r},{f("AD")},"已收讫"),2),'
                       f'ROUND(SUMIFS({f("AJ")},{f("F")},$B{r},{f("X")},"已付清")'
                       f'+SUMIFS({o("CX")},{o("CV")},$B{r},{o("CY")},"已付清"),2)))')
        for c in ('G',):
            ws[f'{c}{r}'] = f'=IF($B{r}="","",ROUND(N($D{r})-N($E{r}),2))'
    a = ws['A205'].value or ''
    if '原料出库' not in a:
        ws['A205'] = a.replace('公司购买果品款＝应付抵扣', '公司购买果品款＝应付抵扣，含原料出库「公司购买」')


# ════════════════════════════════════════════════════════════════════════════
# ⑨ 新表【公司购买接口】（K5）
# ════════════════════════════════════════════════════════════════════════════
def buy_iface(wb):
    ws = wb.create_sheet('公司购买接口')
    ws.sheet_properties.tabColor = 'C55A11'
    ws.sheet_view.showGridLines = False
    A = '_自动清单!'
    cell(ws, 'A1', '笔数', F_LBL, FL_LBL, AC)
    cell(ws, 'B1', f'={A}$CP$1', F_AUTOB, FL_AUTO, AC, fmt='"共 "0" 笔"')
    ws.merge_cells('C1:N1')
    cell(ws, 'C1', '公 司 购 买 接 口（原料出库「公司购买」逐笔 · 全自动 · 供《03 财务账套》取数，请勿改动结构）',
         F_TITLE, FL_TITLE, AC, border=NOB)
    ws.row_dimensions[1].height = 30
    ws.merge_cells('A2:N2')
    cell(ws, 'A2', '★ 只收【原料出库明细】出库原因＝「公司购买」的行，从第 4 行起紧凑排列，不带日期筛选（全部日期）；'
                   '没填采购单价的也收，金额记 0。B1＝笔数。《03》按固定位置引用本表 A3:N403，请勿插入/删除行列、改表名。',
         Font(name='微软雅黑', size=9, color='808080'), FL_NONE, AL, border=NOB)
    ws.row_dimensions[2].height = 32
    heads = ['序号', '日期', '卖方\n（原料货主）', '品种', '等级', '数量', '单位', '采购单价\n（元/件）', '采购金额\n（应付）',
             '付款状态', '购买方/去向', '单号', '原表行号', '备注']
    for i, t in enumerate(heads):
        cell(ws, f'{L(i + 1)}3', t, Font(name='微软雅黑', size=10, bold=True, color='FFFFFF'), FL_HDR, AC)
    ws.row_dimensions[3].height = 32
    for r in range(4, GM_END + 1):
        m = f'$M{r}-3'
        o = lambda col: V(OUT_S, col, OUT_END, m)
        x = lambda col: IX(OUT_S, col, OUT_END, m)
        f = {
            'M': f'=IF(ROW()-3>$B$1,"",MATCH(ROW()-3,{A}$CQ$4:$CQ${OUT_END},0)+3)',
            'A': f'=IF($M{r}="","",ROW()-3)',
            'B': f'=IF($M{r}="","",{o("B")})',
            'C': f'=IF($M{r}="","",{o("F")})',
            'D': f'=IF($M{r}="","",{o("D")})',
            'E': f'=IF($M{r}="","",{o("E")})',
            'F': f'=IF($M{r}="","",N({x("H")}))',
            'G': f'=IF($M{r}="","",{o("I")})',
            'H': f'=IF($M{r}="","",{NV(OUT_S, "U", OUT_END, m)})',
            'I': f'=IF($M{r}="","",N({IX("_自动清单", "CX", OUT_END, m)}))',
            'J': f'=IF($M{r}="","",{o("W")})',
            'K': f'=IF($M{r}="","",{o("X")})',
            'L': f'=IF($M{r}="","",{o("C")})',
            'N': f'=IF($M{r}="","",{o("T")}&IF(N({x("U")})=0,IF({x("T")}=0,"","；")&"未填采购单价",""))',
        }
        for c, v in f.items():
            fmt = {'B': DATE, 'F': QTY, 'H': MONEY_FMT, 'I': MONEY_FMT, 'M': '0'}.get(c)
            cell(ws, f'{c}{r}', v, Font(name='宋体', size=10, color='006100'), FL_AUTO,
                 AL if c in 'N' else AC, fmt=fmt)
    for c, w in {'A': 6, 'B': 11, 'C': 12, 'D': 7, 'E': 11, 'F': 8, 'G': 6, 'H': 10, 'I': 11, 'J': 10,
                 'K': 13, 'L': 12, 'M': 8, 'N': 28}.items():
        ws.column_dimensions[c].width = w
    ws.freeze_panes = 'C4'
    ws.print_title_rows = '3:3'


# ════════════════════════════════════════════════════════════════════════════
# ⑪ 主页、使用说明
# ════════════════════════════════════════════════════════════════════════════
def home(wb):
    ws = wb['主页']
    ws['C13'] = '按货主＋品种＋等级：入库、出库、结余数量与重量，平均重量、在库成品（含损耗）、损耗、损耗率'
    ws['C16'] = '★果然鲜销售/商务/福利/送礼逐笔明细，下面接 一、按销售类型；二、按货主；三、按购买方 三张汇总'
    ws['B39'] = '=HYPERLINK("#\'果然鲜销售明细\'!A478","果然鲜销售明细·按购买方")'
    ws['C39'] = '★按购买方汇总（原【果然鲜销售汇总】已并进【果然鲜销售明细】第 478 行起）：卖了多少、收了多少、还欠多少'
    ws['C40'] = '供《03 财务账套》跨文件取：结余数量/重量、平均重量、在库成品、损耗（不随日期筛选）'
    r = 41
    for c in 'ABCD':
        style_from(ws[f'{c}{r}'], ws[f'{c}40'])
    ws[f'A{r}'] = 5
    ws[f'B{r}'] = '=HYPERLINK("#\'公司购买接口\'!A1","公司购买接口")'
    ws[f'C{r}'] = '★原料出库「公司购买」逐笔（卖方、单价、金额、付款状态、去向），供《03》往来取应付'
    ws[f'D{r}'] = '自动'
    ws.row_dimensions[r].height = ws.row_dimensions[40].height


def manual(wb):
    ws = wb['使用说明']
    ws['A77'] = ('2、新增【果然鲜销售明细】：自动列出销售类型是 果然鲜销售／果然鲜商务／果然鲜福利／果然鲜送礼 的每一笔，'
                 '含货主、品种、等级、数量、采购单价与金额、销售单价与金额、运费，下面有「一、按销售类型」「二、按货主/客户」'
                 '「三、按购买方」三张汇总（10/6 起原【果然鲜销售汇总】并进来，作为第三段）。')
    ws['B39'] = ('原料入库明细预置 5,000 行公式（第 4～5003 行），其他明细表 3,000 行（其他代存 800 行）。'
                 '★ 新的一笔一律接在最后一行下面录，不要在中间插行。')
    ws['B40'] = '超过这些行数，或【对账明细接口】的业务笔数超过 6,000 笔，请联系维护人员用脚本扩容（《04》要一起扩）。'
    items = [
        ('★ 本 次 更 新（2026-10-06）', True),
        ('1、明细表别在中间插行（9 月插的两行让【对账明细接口】漏了最后一笔，现已修好：自动清单改成按位置取数）；'
         '原料入库明细扩到 5,000 行，【对账明细接口】扩到 6,000 行。', False),
        ('2、【库存结余】新增「平均重量KG/件」（全部入库总重÷入库数量），出库总重＝出库数量×平均重量；新增「在库成品吨位（含损耗）」'
         '＝出库加工、二次加工的数量×平均重量－成品出库总重，「损耗」原料出完后才算，「损耗率」＝损耗÷加工出库重量。'
         '成品出库里才有的组合（扎伤、货主＝果然鲜）也列进来了。', False),
        ('3、【原料出库明细】新增 U～X：采购单价（元/件）、采购金额、付款状态、购买方/去向。公司在原料阶段买的货价填这里；'
         '直接买货主加工好的成品，才填【成品出库明细】的采购单价——同一批货别两头都填。', False),
        ('4、【果然鲜采购明细】把原料出库「公司购买」的行也列进来（P 列标来源，没填单价的也列）；新表【公司购买接口】给《03》用。', False),
        ('5、【库存等级接口】补回「平均重量」，新增在库成品、成品出库总重、损耗、加工出库重量，不随日期筛选。', False),
    ]
    r0 = 82
    for i, (t, head) in enumerate(items):
        r = r0 + i
        ws.merge_cells(f'A{r}:H{r}')
        ws[f'A{r}'] = t
        style_from(ws[f'A{r}'], ws['A74'] if head else ws['A75'])
        ws.row_dimensions[r].height = ws.row_dimensions[74 if head else 75].height or (22 if head else 36)
        if not head:
            ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 0, 36)


# ════════════════════════════════════════════════════════════════════════════
def apply(wb):
    widen_refs(wb, skip={'_自动清单', '对账明细接口', '库存结余', '库存等级接口', '果然鲜采购明细', '果然鲜销售明细',
                         '果然鲜销售汇总'})
    out_sheet(wb)
    in_sheet(wb)
    auto_list(wb)
    stock(wb)
    stock_total(wb)
    stock_check(wb)
    stock_iface(wb)
    recon(wb)
    purchase(wb)
    sales(wb)
    finance(wb)
    buy_iface(wb)
    home(wb)
    manual(wb)


def post(path):
    """存盘后的 XML 层处理（只动本册）：
       ① 核对定义名称：localSheetId 指向的表名与名称里的表名一致、没有引用已删的【果然鲜销售汇总】、
          《01》本来没有跨文件链接，存盘后也不应多出来——不对就报错；
       ② 工作表 XML 里 openpyxl 写成 &#21407; 这种字符引用的中文改回 UTF-8 原字（同一个字，文件小一半多）；
       ③ 【对账明细接口】【_自动清单】同一列上下行写法一样的公式，改成 Excel/WPS 自己存盘时用的「共享公式」
          （只在第一格写全文，下面各格引用它），算出来完全一样，文件小很多、WPS 打开快。"""
    z = zipfile.ZipFile(path)
    wbx = z.read('xl/workbook.xml').decode('utf8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf8')
    sheets = re.findall(r'<sheet [^>]*name="([^"]+)"', wbx)
    bad = []
    for m in re.finditer(r'<definedName ([^>]*)>([^<]*)</definedName>', wbx):
        attrs, val = m.group(1), m.group(2)
        lid = re.search(r'localSheetId="(\d+)"', attrs)
        if '果然鲜销售汇总' in val:
            bad.append(('引用已删表', attrs, val))
        if lid:
            want = sheets[int(lid.group(1))]
            ref_sheet = val.split('!')[0].strip("'")
            if ref_sheet != want:
                bad.append(('localSheetId 错位', attrs, val, want))
    if '果然鲜销售汇总' in sheets:
        bad.append(('果然鲜销售汇总还在',))
    if 'externalLink' in rels:
        bad.append(('多出了跨文件链接',))
    if bad:
        z.close()
        raise SystemExit(f'《01》存盘后核对不过：{bad}')
    # 表名 → 部件路径
    rid = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship [^>]*?Id="([^"]+)"[^>]*?Target="([^"]+)"', rels)}
    rid.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship [^>]*?Target="([^"]+)"[^>]*?Id="([^"]+)"', rels)})
    part = {}
    for m in re.finditer(r'<sheet [^>]*?name="([^"]+)"[^>]*?r:id="([^"]+)"', wbx):
        t = rid[m.group(2)].lstrip('/')
        part['xl/' + t if not t.startswith('xl/') else t] = m.group(1)
    tmp = path + '.tmp'
    zo = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in z.infolist():
        data = z.read(it.filename)
        if it.filename in part:
            s = _utf8_refs(data.decode('utf8'))
            if part[it.filename] in ('对账明细接口', '_自动清单'):
                s = _share(s)
            data = s.encode('utf8')
        zo.writestr(it, data)
    zo.close()
    z.close()
    import shutil
    shutil.move(tmp, path)
    return len(sheets)


def _utf8_refs(s):
    """&#21407; / &#x539f; → 原（只换 127 以上的字符；< & > 这些 ASCII 转义原样留着）"""
    def one(m):
        n = int(m.group(1)) if m.group(1) else int(m.group(2), 16)
        return chr(n) if n > 127 else m.group(0)
    return re.sub(r'&#(?:(\d+)|x([0-9a-fA-F]+));', one, s)


_STR = re.compile(r'"(?:[^"]|"")*"')
_REF = re.compile(r'(?<![A-Za-z0-9_$.])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_])')
_FCELL = re.compile(r'<c r="([A-Z]{1,3})(\d+)"([^>]*)><f>([^<]*)</f>')


def _norm(f, row):
    """把公式里「相对行号」换成相对本行的偏移；同一列里两格的 _norm 相同＝可以共享"""
    out, i = [], 0
    for m in _STR.finditer(f):
        out.append(_REF.sub(lambda x: x.group(0) if x.group(3) else
                            f'{x.group(1)}{x.group(2)}R[{int(x.group(4)) - row}]', f[i:m.start()]))
        out.append(m.group(0))
        i = m.end()
    out.append(_REF.sub(lambda x: x.group(0) if x.group(3) else
                        f'{x.group(1)}{x.group(2)}R[{int(x.group(4)) - row}]', f[i:]))
    return ''.join(out)


def _share(s):
    groups = {}          # 列 → [首行, 末行, 规范式, 首格公式]
    done = []
    for m in _FCELL.finditer(s):
        col, row, f = m.group(1), int(m.group(2)), m.group(4)
        nf = _norm(f, row)
        g = groups.get(col)
        if g and g[1] == row - 1 and g[2] == nf:
            g[1] = row
        else:
            if g:
                done.append((col, g[0], g[1]))
            groups[col] = [row, row, nf]
    done += [(c, g[0], g[1]) for c, g in groups.items()]
    role = {}
    si = 0
    for col, r0, r1 in done:
        if r1 - r0 < 1:
            continue
        role[(col, r0)] = ('M', si, f'{col}{r0}:{col}{r1}')
        for r in range(r0 + 1, r1 + 1):
            role[(col, r)] = ('S', si, None)
        si += 1

    def rep(m):
        col, row = m.group(1), int(m.group(2))
        x = role.get((col, row))
        if not x:
            return m.group(0)
        head = f'<c r="{col}{row}"{m.group(3)}>'
        if x[0] == 'M':
            return f'{head}<f t="shared" ref="{x[2]}" si="{x[1]}">{m.group(4)}</f>'
        return f'{head}<f t="shared" si="{x[1]}"/>'
    return _FCELL.sub(rep, s)
