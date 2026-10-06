# -*- coding: utf-8 -*-
"""10/6 这一轮（A050 1006）《04 综合查询对账单》的改动。底稿：参考/1006上传/04_综合查询对账单模板.xlsx。

用法：wb = openpyxl.load_workbook(底稿)；apply(wb)；wb.save(成品)；post(成品)。
对底稿重复跑，结果一样（所有公式、版式都按固定规则从底稿重写，不依赖上一次的结果）。

这一轮改了什么（对应 A050 1006 实施规格第 5 节，接口 K1 / K2 / K3 的读取方）：

① 扩容（K1、K2）：【对接源_01水果】逐格对应《01》【对账明细接口】第 4～6003 行（1,200 → 6,000 行）；
   【对接源_02物料】对应《02》【对账明细接口】第 4～3003 行（1,200 → 3,000 行）。
   【_自动清单】整张重写：第 4～6003 行对 01、6004～9003 行对 02；逐行取数改成按位置
   INDEX(对接源!$C$4:$C$6003,ROW()-3)；各对账单里「<=1200 / -1200」的分段阈值改成 6000，
   所有 $1203 / $2403 的区间、计数、MATCH 范围一起放大；紧凑清单和对账单的 MATCH 辅助列都加了
   「超过笔数就不查」护栏（不然每格都要把 9,000 行扫一遍）。
② 各对账单明细段扩容，每段标题后面显示「共 N 笔」，超过显示行数时整条标红并提示「下面只列前 M 笔，请缩短起止日期」：
   筐子对账单 出库/退回/随筐入库 100/60/60 → 300/300/300；自备筐对账单 入库/出库 100/100 → 300/300；
   水果对账单 原料入库/原料出库/成品出库 120/80/100 → 400/200/300；物料与筐子销售对账单 200 → 400；
   果然鲜销售对账单 200 → 400；公司购买果品对账单 150 → 300。
   段与段之间的标题、表头、合并格、打印区都跟着下移。各段合计本来就是对全部数据 SUMIFS，不受显示行数限制；
   水果对账单「入库/出库等级汇总」原来只按下面显示出来的明细求和，改成直接对全部数据按等级 SUMIFS，
   合计行改成全部合计；果然鲜销售对账单底部合计行也改成取右侧全量汇总。
③ 筐子库存汇总（K3）：《02》【筐子库存汇总】第 4 行插了「合计」，数据改在 5～64 行；本表 A4:M63 改成
   逐格读《02》A5:M64（本表第 64 行自己的合计不变），表头文字跟《02》对齐。
④ 公司购买果品对账单：右侧「采购总数量」按单位拆成「采购数量（筐等）」「采购数量（箱）」两行
   （原料出库按筐、成品出库按箱，两者不能相加）；说明改成也包括《01》原料出库「公司购买」的行。
⑤ 物料与筐子销售对账单：明细和汇总都排除「包装物料－采购入库」（新顺发泡网的进货不再当成要收的钱），
   一并排除「包装物料－期初库存」（建账库存，不是销售）。
⑥ 应收应付汇总：成品出库的应收改按购买方算（《01》接口 X＝购买方，没填购买方的用货主），不再记在货主头上；
   新增「公司购买果品（应付）」列＝按采购货主汇总应付合计（负数）、「物料/筐子采购（应付）」列＝《02》
   包装物料和周转筐「采购入库」的金额（负数，原来混在包装物料、周转筐列里当成应收）；已收/已付里
   采购已付款记负数；往来单位名单并入购买方。G 列表头注明「使用费，不含押金」（《02》接口次果段只给使用费）；
   A2 说明末尾提示已收/已付只是《01》《02》上标的收付款状态，欠多少以《03》往来对账单明细为准。
   列宽用 set_col_widths 拆列组后再设，<col> 不重叠。本模块不设冻结窗格（各表冻结沿用底稿）。
⑦ 主页、使用说明简短更新；对接源表头显示「已取到 N/6000 笔」。

post(path)：存盘后把【_自动清单】和各对账单辅助列里上下行写法一样的公式改成「共享公式」
   （Excel/WPS 自己存盘也是这样，算出来完全一样，文件小很多），工作表 XML 里 openpyxl 写成 &#xxxx; 的中文改回原字。
"""
import copy
import re
import zipfile
import shutil

from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.dimensions import ColumnDimension

S1, S2, AUTO = '对接源_01水果', '对接源_02物料', '_自动清单'
N1, N2 = 6000, 3000                 # K1：01 对账明细接口 6,000 行；K2：02 对账明细接口 3,000 行
E1, E2 = 3 + N1, 3 + N2             # 6003 / 3003
AE = E1 + N2                        # 9003：_自动清单 第 4～6003 行对 01，6004～9003 行对 02
OLD_END, OLD_N, OLD_AE = 1203, 1200, 2403


def R1(c):
    return f'{S1}!${c}$4:${c}${E1}'


def R2(c):
    return f'{S2}!${c}$4:${c}${E2}'


# ════════════════════════════════════════════════════════════════════════════
# 公式小工具
# ════════════════════════════════════════════════════════════════════════════
_STR = re.compile(r'"(?:[^"]|"")*"')
_REF = re.compile(r'(?<![A-Za-z0-9_$.\]])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_!])')


def _outside_strings(f, fn):
    out, i = [], 0
    for m in _STR.finditer(f):
        out.append(fn(f[i:m.start()]))
        out.append(m.group(0))
        i = m.end()
    out.append(fn(f[i:]))
    return ''.join(out)


def shift_rows(f, d):
    """像 Excel 复制粘贴一样：公式里不带 $ 的行号整体平移 d 行（字符串常量不动）"""
    if not d or not isinstance(f, str) or not f.startswith('='):
        return f

    def one(seg):
        return _REF.sub(lambda m: m.group(0) if m.group(3) else
                        f'{m.group(1)}{m.group(2)}{int(m.group(4)) + d}', seg)
    return _outside_strings(f, one)


_RX = [
    (re.compile(r"((?:'?对接源_01水果'?)!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)1203(?!\d)"), rf'\g<1>{E1}'),
    (re.compile(r"((?:'?对接源_02物料'?)!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)1203(?!\d)"), rf'\g<1>{E2}'),
    (re.compile(r"((?:'?_自动清单'?)!\$[A-Z]{1,3}\$4:\$[A-Z]{1,3}\$)2403(?!\d)"), rf'\g<1>{AE}'),
    (re.compile(r"(\$A[EFG]\$?\d+)<=1200(?!\d)"), rf'\g<1><={N1}'),
    (re.compile(r"(\$A[EFG]\$?\d+)-1200(?!\d)"), rf'\g<1>-{N1}'),
]


def shift_abs(f, lo, hi, d):
    """公式里带 $ 的行号落在 lo～hi 的，平移 d 行（搬底部合计区时用）"""
    if not d or not isinstance(f, str) or not f.startswith('='):
        return f

    def one(seg):
        return _REF.sub(lambda m: f'{m.group(1)}{m.group(2)}{m.group(3)}{int(m.group(4)) + d}'
                        if m.group(3) and lo <= int(m.group(4)) <= hi else m.group(0), seg)
    return _outside_strings(f, one)


def retarget(f):
    """旧区间 → 新区间：对接源 $1203 → $6003/$3003，_自动清单 $2403 → $9003，分段阈值 1200 → 6000"""
    if not isinstance(f, str) or not f.startswith('='):
        return f

    def one(seg):
        for rx, rep in _RX:
            seg = rx.sub(rep, seg)
        return seg
    return _outside_strings(f, one)


def put(ws, ref, value=None, style=None):
    c = ws[ref]
    c.value = value
    if style is not None:
        c._style = copy.copy(style)
    return c


def drop(ws, r, c):
    """整格删掉（值和样式），比写 None 干净"""
    ws._cells.pop((r, c), None)


def set_col_widths(ws, widths):
    """按列设宽度，不让 <col> 定义互相重叠：先把覆盖到这些列的列组（如 C:J、M:XFD）拆开，
    拆出来的每一段照抄原组的样式/隐藏等属性，只改指定列的宽度；结果按列号连续、互不重叠、最大到 XFD（16384）。"""
    cd = ws.column_dimensions
    owner = {}
    for k, d in sorted(cd.items(), key=lambda kv: kv[1].min or CI(kv[0])):
        lo = d.min or CI(k)
        hi = min(d.max or lo, 16384)
        for c in range(lo, hi + 1):
            owner[c] = d                      # 万一原来就有重叠：按起始列排序，后一条覆盖前一条
    want = {CI(c): w for c, w in widths.items()}
    for c in want:
        owner.setdefault(c, None)
    runs = []
    for c in sorted(owner):
        key = (id(owner[c]), want.get(c))
        if runs and runs[-1][1] == c - 1 and runs[-1][2] == key:
            runs[-1][1] = c
        else:
            runs.append([c, c, key, owner[c], want.get(c)])
    cd.clear()
    for lo, hi, _, d, w in runs:
        nd = copy.copy(d) if d is not None else ColumnDimension(ws, index=L(lo))
        nd.index, nd.min, nd.max = L(lo), lo, hi
        if w is not None:
            nd.width = w
        cd[L(lo)] = nd


# ════════════════════════════════════════════════════════════════════════════
# ① 两张对接源：6,000 / 3,000 行逐格跨文件取数
# ════════════════════════════════════════════════════════════════════════════
_XREF = re.compile(r"^=IFERROR\(IF\((\[\d+\])对账明细接口!\$([A-Z]{1,3})(\$?)4=0,")


def src_sheet(wb, name, end, cap):
    ws = wb[name]
    link, cmap, const = None, {}, {}
    for c in range(2, ws.max_column + 1):
        v = ws.cell(row=4, column=c).value
        if not isinstance(v, str):
            continue
        m = _XREF.match(v)
        if m:
            link = m.group(1)
            cmap[L(c)] = (m.group(2), m.group(3))      # 底稿 B～M 写 $B$4，S～AJ 写 $N4，照原样
        elif v == '=""':
            const[L(c)] = v
    assert link and len(cmap) >= 18, (name, link, cmap)
    tpl = {L(c): copy.copy(ws.cell(row=4, column=c)._style) for c in range(1, ws.max_column + 1)
           if ws.cell(row=4, column=c).has_style}
    h = ws.row_dimensions[4].height
    for r in range(4, end + 1):
        for c, st in tpl.items():
            if r > OLD_END or r == 4:
                ws[f'{c}{r}']._style = copy.copy(st)
        ws[f'A{r}'] = r
        for c, (sc, d) in cmap.items():
            x = f'{link}对账明细接口!${sc}{d}{r}'
            ws[f'{c}{r}'] = f'=IFERROR(IF({x}=0,"",{x}),"")'
        for c, v in const.items():
            ws[f'{c}{r}'] = v
        ws[f'N{r}'] = f'=IF($C{r}="",0,1)'
        ws[f'O{r}'] = f'=N(O{r - 1})+$N{r}'
        if h:
            ws.row_dimensions[r].height = h
    ws['R1'] = f'=MAX($O$4:$O${end})'
    ws['R3'] = f'=COUNTA($C$4:$C${end})'
    book = '01 水果进销存台账' if name == S1 else '02 物料与周转物台账'
    ws['A2'] = (f'=IF($R$4=1,"✔ 链接正常，已取到 "&$R$1&"/{cap} 笔业务"&IF(N($R$1)>={cap},"（已满，请联系维护人员扩容）","")&'
                f'"　｜　本表第 4～{end} 行 1:1 对应《{book}》的【对账明细接口】第 4～{end} 行，三个文件必须放在同一个文件夹，'
                f'文件名也不要改","★★ 链接没接上！→ 请看本表右边 Q6 格的修复说明（往右拉一屏）")')


# ════════════════════════════════════════════════════════════════════════════
# ① _自动清单：整张按新行数重写
# ════════════════════════════════════════════════════════════════════════════
def auto_list(wb):
    ws = wb[AUTO]
    # 先清掉第 4 行以下的全部旧公式
    for (r, c) in [k for k in ws._cells if k[0] >= 4]:
        del ws._cells[(r, c)]

    def X(r, col):
        if r <= E1:
            return f'INDEX({R1(col)},ROW()-3)'
        return f'INDEX({R2(col)},ROW()-{E1})'

    def flag(r, sel, cond):
        """sel：K/L/M/N（各对账单的客户＋日期命中）；cond：本类业务的条件"""
        return f'=IF(AND(${sel}{r}=1,{cond}),1,0)'

    pairs = ['P', 'R', 'T', 'V', 'X', 'Z', 'AB', 'AD', 'AF', 'AH']       # 命中列；累计列＝右边一列
    for p in pairs:
        cum = L(CI(p) + 1)
        ws[f'{p}1'] = f'=MAX({cum}$4:{cum}${AE})'
    ws['AM1'] = f'=MAX(AM$4:AM${AE})'
    ws['AQ1'] = f'=N(AP${E1})'
    ws['AR1'] = f'=MAX(AS$4:AS${AE})'
    ws['AN3'] = '应收单位（成品出库：购买方，没填用货主）'
    ws['AQ3'] = '应收单位(紧凑)'
    ws['AU3'] = '往来单位∪应收单位'
    ws['AV3'] = '首次'
    ws['AW3'] = '累计'
    ws['AX3'] = '应收应付汇总单位(紧凑)'
    ws['AX1'] = '=MAX(AW$4:AW$503)'
    ws['AU1'] = '【应收应付汇总】用：E 列往来单位＋AQ 列成品出库购买方，去重'
    ws['A1'] = ('辅助表：K/L/M/N/AL 第1~3行＝各对账单的 客户/开始日期/结束日期，请勿改动。'
                f'第 4～{E1} 行对应【对接源_01水果】，第 {E1 + 1}～{AE} 行对应【对接源_02物料】')

    for r in range(4, AE + 1):
        one = r <= E1
        C, B, D, Xc = X(r, 'C'), X(r, 'B'), X(r, 'D'), X(r, 'X')
        ws[f'A{r}'] = f'=IF({C}="","",{C})'
        ws[f'B{r}'] = f'=IF(A{r}="",0,IF(MATCH(A{r},$A$4:$A${AE},0)=ROW()-3,1,0))'
        ws[f'C{r}'] = f'=N(C{r - 1})+B{r}'
        for s in ('K', 'L', 'M', 'N'):
            ws[f'{s}{r}'] = (f'=IF(OR(${s}$1="",{C}="",{C}<>${s}$1,{B}=""),0,'
                             f'IF(OR(N({B})<${s}$2,N({B})>${s}$3),0,1))')
        ws[f'P{r}'] = flag(r, 'K', f'{Xc}="公司筐",N({X(r, "T")})<>0')
        ws[f'R{r}'] = flag(r, 'K', f'{Xc}="公司筐",N({X(r, "U")})<>0')
        ws[f'T{r}'] = flag(r, 'K', f'{Xc}="公司筐",N({X(r, "V")})<>0')
        ws[f'V{r}'] = flag(r, 'L', f'{Xc}="自备筐",N({X(r, "V")})<>0')
        ws[f'X{r}'] = flag(r, 'L', f'{Xc}="自备筐",N({X(r, "T")})<>0')
        ws[f'Z{r}'] = flag(r, 'M', f'{D}="原料入库"')
        ws[f'AB{r}'] = flag(r, 'M', f'LEFT({D},4)="原料出库"')
        ws[f'AD{r}'] = flag(r, 'M', f'LEFT({D},4)="成品出库"')
        ws[f'AF{r}'] = flag(r, 'N', f'LEFT({D},4)="包装物料",{D}<>"包装物料－采购入库",{D}<>"包装物料－期初库存",'
                                    f'N({X(r, "K")})<>0')
        ws[f'AH{r}'] = flag(r, 'N', f'{D}="周转筐－销售出库",N({X(r, "K")})<>0')
        for p in pairs:
            cum = L(CI(p) + 1)
            ws[f'{cum}{r}'] = f'=N({cum}{r - 1})+${p}{r}'
        if one:
            Y, Bx = X(r, 'Y'), B
            ws[f'AL{r}'] = (f'=IF(OR($AL$1="",{Y}="",{Y}<>$AL$1,{Bx}="",N({X(r, "AJ")})<=0),0,'
                            f'IF(OR(N({Bx})<$AL$2,N({Bx})>$AL$3),0,1))')
            ws[f'AN{r}'] = (f'=IF(LEFT({D},4)<>"成品出库","",IF({X(r, "AC")}<>"",{X(r, "AC")},{C}))')
            ws[f'AO{r}'] = f'=IF(AN{r}="",0,IF(MATCH(AN{r},$AN$4:$AN${E1},0)=ROW()-3,1,0))'
            ws[f'AP{r}'] = f'=N(AP{r - 1})+$AO{r}'
        else:
            ws[f'AL{r}'] = 0
        ws[f'AM{r}'] = f'=N(AM{r - 1})+$AL{r}'
        ws[f'AR{r}'] = f'=IF(OR($AF{r}=1,$AH{r}=1),1,0)'
        ws[f'AS{r}'] = f'=N(AS{r - 1})+$AR{r}'
    for r in range(4, 304):
        ws[f'E{r}'] = (f'=IF(ROW()-3>$C${AE},"",IFERROR(INDEX($A$4:$A${AE},MATCH(ROW()-3,$C$4:$C${AE},0)),""))')
    for r in range(4, 204):
        ws[f'AQ{r}'] = (f'=IF(ROW()-3>$AQ$1,"",IFERROR(INDEX($AN$4:$AN${E1},MATCH(ROW()-3,$AP$4:$AP${E1},0)),""))')
    # 应收应付汇总用的单位：E 列（两册往来单位）＋ AQ 列（成品出库购买方），去重
    for r in range(4, 504):
        src = f'INDEX($E$4:$E$303,ROW()-3)' if r <= 303 else f'INDEX($AQ$4:$AQ$203,ROW()-303)'
        ws[f'AU{r}'] = f'=IF({src}="","",{src})'
        ws[f'AV{r}'] = f'=IF(AU{r}="",0,IF(MATCH(AU{r},$AU$4:$AU$503,0)=ROW()-3,1,0))'
        ws[f'AW{r}'] = f'=N(AW{r - 1})+AV{r}'
    for r in range(4, 304):
        ws[f'AX{r}'] = (f'=IF(ROW()-3>$AX$1,"",IFERROR(INDEX($AU$4:$AU$503,MATCH(ROW()-3,$AW$4:$AW$503,0)),""))')


# ════════════════════════════════════════════════════════════════════════════
# ② 对账单明细段扩容（按段搬版式，不用插行）
# ════════════════════════════════════════════════════════════════════════════
RED_FILL = PatternFill(fill_type='solid', start_color='FFC00000', end_color='FFC00000')
RED_FONT = Font(color='FFFFFFFF', bold=True)

# 段：(标题行, 表头行, 首数据行, 末数据行, 本表辅助列, _自动清单累计列, _自动清单计数格)
LAYOUT = {
    '自备筐对账单': dict(cols='ABCDEFGHI', segs=[(4, 5, 6, 105, 'AE', 'W', 'V1'), (107, 108, 109, 208, 'AF', 'Y', 'X1')],
                    sizes=[300, 300], pcols='M'),
    '筐子对账单': dict(cols='ABCDEFG', segs=[(4, 5, 6, 105, 'AE', 'Q', 'P1'), (107, 108, 109, 168, 'AF', 'S', 'R1'),
                                            (170, 171, 172, 231, 'AG', 'U', 'T1')], sizes=[300, 300, 300], pcols='K'),
    '物料与筐子销售对账单': dict(cols='ABCDEFGHIJ', segs=[(4, 5, 6, 205, 'AE', 'AS', 'AR1')], sizes=[400], pcols='J'),
    '水果对账单': dict(cols='ABCDEFGHIJ', segs=[(4, 5, 6, 125, 'AE', 'AA', 'Z1'), (127, 128, 129, 208, 'AF', 'AC', 'AB1'),
                                               (210, 211, 212, 311, 'AG', 'AE', 'AD1')], sizes=[400, 200, 300], pcols='N',
                  extra=['AI', 'AJ', 'AK', 'AL', 'AM', 'AN']),
    '公司购买果品对账单': dict(cols='ABCDEFGHIJKL', segs=[(4, 5, 6, 155, 'AE', 'AM', 'AM1')], sizes=[300], pcols='N'),
}


def count_label(text, cnt, n):
    """标题后面接「共 N 笔」；超过显示行数时提示只列前 n 笔"""
    return f'="{text}　共 "&{cnt}&" 笔"&IF({cnt}>{n},"，下面只列前 {n} 笔，请缩短起止日期","")'


def add_overflow_cf(ws, rng, cnt, n):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{cnt}>{n}'], font=RED_FONT, fill=RED_FILL,
                                                   stopIfTrue=True))


def relayout(wb, name, spec):
    ws = wb[name]
    cols = list(spec['cols'])
    segs, sizes = spec['segs'], spec['sizes']
    extra = spec.get('extra', [])
    helpers = [s[4] for s in segs]
    allcols = cols + helpers + extra
    old_last = segs[-1][3]
    # 新位置
    new, t = [], 4
    for n in sizes:
        new.append((t, t + 1, t + 2, t + 1 + n))
        t = t + 3 + n
    new_last = new[-1][3]
    # 快照
    snap = {}
    for r in range(4, old_last + 1):
        for c in allcols:
            x = ws._cells.get((r, CI(c)))
            snap[(r, c)] = (x.value, copy.copy(x._style) if x is not None and x.has_style else None) if x else (None, None)
    hts = {r: ws.row_dimensions[r].height for r in range(4, old_last + 1)}
    c0, c1 = CI(cols[0]), CI(cols[-1])
    moved = []
    for m in sorted(ws.merged_cells.ranges, key=lambda m: (m.min_row, m.min_col)):
        if m.min_row >= 4 and m.max_row <= old_last and m.min_col >= c0 and m.max_col <= c1:
            moved.append((m.min_row, m.min_col, m.max_row, m.max_col))
            ws.unmerge_cells(str(m))
    # 清场
    for r in range(4, max(old_last, new_last) + 2):
        for c in allcols:
            drop(ws, r, CI(c))
        ws.row_dimensions[r].height = None
    # 逐段写
    for (ot, oh, od0, od1, hc, cum, cnt), (nt, nh, nd0, nd1), n in zip(segs, new, sizes):
        for orow, nrow in ((ot, nt), (oh, nh)):
            for c in cols:
                v, st = snap[(orow, c)]
                if v is not None or st is not None:
                    put(ws, f'{c}{nrow}', shift_rows(retarget(v), nrow - orow), st)
            ws.row_dimensions[nrow].height = hts[orow]
        k = nd0 - 1
        for j in range(n):
            r = nd0 + j
            tpl = od0 + (j % 2)
            for c in cols:
                v, _ = snap[(od0, c)]
                _, st = snap[(tpl, c)]
                put(ws, f'{c}{r}', shift_rows(retarget(v), r - od0), st)
            hst = snap[(od0, hc)][1]
            put(ws, f'{hc}{r}', f'=IF(ROW()-{k}>${hc}${nt},"",IFERROR(MATCH(ROW()-{k},{AUTO}!${cum}$4:${cum}${AE},0),""))',
                hst)
            ws.row_dimensions[r].height = hts[od0]
        # 计数格（本表辅助列的标题行）＋ 标题文字
        cref = f'${hc}${nt}'
        mc = re.match(r'([A-Z]+)(\d+)$', cnt)
        put(ws, f'{hc}{nt}', f'=N({AUTO}!${mc.group(1)}${mc.group(2)})', snap[(od0, hc)][1])
        lab = snap[(ot, cols[0])][0]
        base = lab if isinstance(lab, str) and not lab.startswith('=') else str(lab)
        a = ws[f'{cols[0]}{nt}']
        a.value = count_label(base, cref, n)
        al = copy.copy(a.alignment)
        al.wrap_text = True
        a.alignment = al
        ws.row_dimensions[nt].height = max(hts[ot] or 24, 30)
        # 合并格跟着搬
        lab_end = None
        for (r1, cc1, r2, cc2) in moved:
            if r1 == ot:
                ws.merge_cells(start_row=nt, start_column=cc1, end_row=nt + (r2 - r1), end_column=cc2)
                if cc1 == c0:
                    lab_end = cc2
            elif r1 == oh:
                ws.merge_cells(start_row=nh, start_column=cc1, end_row=nh + (r2 - r1), end_column=cc2)
        add_overflow_cf(ws, f'{cols[0]}{nt}:{L(lab_end or c0)}{nt}', cref, n)
    ws.print_area = f'A1:{spec["pcols"]}{new_last + 1}'
    return new


# ── 水果对账单：等级汇总改成对全部数据 SUMIFS；去重辅助列跟着新段位 ───────────────────
def fruit_grades(wb, new):
    ws = wb['水果对账单']
    (t1, h1, a0, a1), (t2, h2, b0, b1), _ = new
    f_in = f'{R1("C")},$B$2,{R1("B")},">="&$P$1,{R1("B")},"<="&$P$2'
    for (d0, d1, ci, cj, ck, l0, l1, crit) in ((a0, a1, 'AI', 'AJ', 'AK', 12, 29, '"原料入库"'),
                                               (b0, b1, 'AL', 'AM', 'AN', 34, 51, '"原料出库*"')):
        for r in range(d0, d1 + 1):
            ws[f'{ci}{r}'] = f'=IF($D{r}="","",$D{r})'
            ws[f'{cj}{r}'] = f'=IF({ci}{r}="",0,IF(MATCH({ci}{r},{ci}${d0}:{ci}${d1},0)=ROW()-{d0 - 1},1,0))'
            ws[f'{ck}{r}'] = f'=N({ck}{r - 1})+{cj}{r}' if r > d0 else f'={cj}{r}'
            for c in (ci, cj, ck):
                ws[f'{c}{r}'].font = Font(name='宋体', size=11, color='FFBFBFBF')
        for r in range(l0, l1 + 1):
            ws[f'L{r}'] = (f'=IF(ROW()-{l0 - 1}>N({ck}${d1}),"",IFERROR(INDEX($D${d0}:$D${d1},'
                           f'MATCH(ROW()-{l0 - 1},{ck}${d0}:{ck}${d1},0)),""))')
            ws[f'M{r}'] = f'=IF($L{r}="","",SUMIFS({R1("H")},{f_in},{R1("D")},{crit},{R1("G")},$L{r}))'
            ws[f'N{r}'] = f'=IF($L{r}="","",ROUND(SUMIFS({R1("J")},{f_in},{R1("D")},{crit},{R1("G")},$L{r}),2))'
        ws[f'M{l1 + 1}'] = f'=ROUND(SUMIFS({R1("H")},{f_in},{R1("D")},{crit}),0)'
        ws[f'N{l1 + 1}'] = f'=ROUND(SUMIFS({R1("J")},{f_in},{R1("D")},{crit}),2)'
    ws['L10'] = '入 库 等 级 汇 总（全部数据）'
    ws['L32'] = '出 库 等 级 汇 总（全部数据）'


# ── 果然鲜销售对账单：单段＋底部合计/签字区，T/U 是本表隐藏辅助列 ─────────────────────
def gx_sales(wb):
    ws = wb['果然鲜销售对账单']
    cols = list('ABCDEFGHIJKLMN')
    D0, D1_OLD, N = 6, 205, 400
    D1 = D0 + N - 1                         # 405
    F_OLD, F_NEW = D1_OLD + 1, D1 + 1       # 合计行 206 → 406
    snap = {}
    for r in list(range(D0, D0 + 2)) + list(range(F_OLD, F_OLD + 5)):
        for c in cols + ['T']:
            x = ws._cells.get((r, CI(c)))
            snap[(r, c)] = (x.value, copy.copy(x._style) if x is not None and x.has_style else None) if x else (None, None)
    hts = {r: ws.row_dimensions[r].height for r in list(range(D0, D0 + 2)) + list(range(F_OLD, F_OLD + 5))}
    for m in list(ws.merged_cells.ranges):
        if m.min_row >= F_OLD:
            ws.unmerge_cells(str(m))
    for r in range(D0, F_NEW + 6):
        for c in cols + ['T']:
            drop(ws, r, CI(c))
        ws.row_dimensions[r].height = None
    for j in range(N):
        r = D0 + j
        tpl = D0 + (j % 2)
        for c in cols:
            v, _ = snap[(D0, c)]
            put(ws, f'{c}{r}', shift_rows(retarget(v), r - D0), snap[(tpl, c)][1])
        put(ws, f'T{r}', f'=IF(ROW()-5>$T$4,"",IFERROR(MATCH(ROW()-5,$U$4:$U${E1},0),""))', snap[(D0, 'T')][1])
        ws.row_dimensions[r].height = hts[D0]
    for k in range(5):
        for c in cols:
            v, st = snap[(F_OLD + k, c)]
            if v is not None or st is not None:
                put(ws, f'{c}{F_NEW + k}', shift_abs(shift_rows(v, F_NEW - F_OLD), F_OLD, F_OLD + 4, F_NEW - F_OLD), st)
        ws.row_dimensions[F_NEW + k].height = hts[F_OLD + k]
    # 底部合计＝右边全量汇总（不受显示 400 行限制）
    for c, src in (('F', 'R3'), ('I', 'R4'), ('J', 'R5'), ('K', 'R6'), ('L', 'R7')):
        ws[f'{c}{F_NEW}'] = f'=ROUND(N(${src[0]}${src[1:]}),2)'
    ws.merge_cells(f'A{F_NEW + 2}:E{F_NEW + 2}')
    # U：逐行命中累计（对接源 4～6003）
    ust = copy.copy(ws['U4']._style) if ws['U4'].has_style else None
    for r in range(4, E1 + 1):
        if r > OLD_END:
            drop(ws, r, CI('U'))
        b = f'INDEX({R1("B")},ROW()-3)'
        put(ws, f'U{r}', f'=N(U{r - 1})+IF(AND($B$2<>"",INDEX({R1("AH")},ROW()-3)=$B$2,{b}>=$U$1,{b}<=$U$2),1,0)', ust)
    ws['T4'] = f'=N($U${E1})'
    ws['T4'].font = Font(name='宋体', size=9, color='FFBFBFBF')
    a = ws['A4']
    a.value = count_label('所选类型未收金额', '$T$4', N)
    al = copy.copy(a.alignment)
    al.wrap_text = True
    a.alignment = al
    ws.row_dimensions[4].height = 30
    add_overflow_cf(ws, 'A4:E4', '$T$4', N)


# ════════════════════════════════════════════════════════════════════════════
# ③ 筐子库存汇总（K3）
# ════════════════════════════════════════════════════════════════════════════
def basket_stock(wb):
    ws = wb['筐子库存汇总']
    rx = re.compile(r"(\[\d+\]筐子库存汇总!\$[A-Z]{1,3}\$)(\d+)")
    n = 0
    for r in range(4, 64):
        for c in range(1, 14):
            x = ws.cell(row=r, column=c)
            if isinstance(x.value, str) and '筐子库存汇总!' in x.value:
                new = rx.sub(lambda m: f'{m.group(1)}{r + 1}' if int(m.group(2)) == r else m.group(0), x.value)
                assert new != x.value, (r, c, x.value)
                x.value = new
                n += 1
    assert n == 60 * 13, n
    for c, t in (('C', '期初库存\n(不受日期限制)'), ('H', '其他入库（调拨等）'), ('J', '客户未回库\n(02册)'),
                 ('K', '客户未回库\n(01水果册·公司筐)')):
        ws[f'{c}3'] = t
    ws['A2'] = ('=IF(N(对接源_02物料!$R$4)=1,"★ 全自动：整表逐格取自《02 物料与周转物台账》【筐子库存汇总】第 5～64 行（那边第 4 行是合计）；'
                '起止日期跟《02》那张表走。公司在库结存＝期初＋采购＋客户退回−发出−销售＋其他入库；公司筐总量＝在库＋客户手上未回的",'
                '"★★ 取不到《02 物料与周转物台账模板.xlsx》的数据，请到【对接源_02物料】按提示修复链接")')


# ════════════════════════════════════════════════════════════════════════════
# ④ 公司购买果品对账单：采购数量按筐 / 箱拆开
# ════════════════════════════════════════════════════════════════════════════
def company_buy(wb):
    ws = wb['公司购买果品对账单']
    old = {r: {c: (ws[f'{c}{r}'].value, copy.copy(ws[f'{c}{r}']._style)) for c in 'NOP'} for r in range(3, 9)}
    for m in list(ws.merged_cells.ranges):
        if m.min_col == CI('N') and m.max_col == CI('O') and 3 <= m.min_row <= 8:
            ws.unmerge_cells(str(m))
    src = {3: 3, 4: 3, 5: 4, 6: 5, 7: 6, 8: 7, 9: 8}
    for r, o in src.items():
        for c in 'NOP':
            v, st = old[o][c]
            put(ws, f'{c}{r}', retarget(v) if c != 'O' else v, st)
        ws.merge_cells(f'N{r}:O{r}')
    base = (f'{R1("Y")},$B$2,{R1("B")},">="&$R$1,{R1("B")},"<="&$R$2,{R1("AJ")},">0"')
    ws['N3'] = '采购数量（筐等）'
    ws['P3'] = f'=ROUND(SUMIFS({R1("H")},{base},{R1("I")},"<>箱"),0)'
    ws['N4'] = '采购数量（箱）'
    ws['P4'] = f'=ROUND(SUMIFS({R1("H")},{base},{R1("I")},"箱"),0)'
    ws['P9'] = '=ROUND(N($P$7)-N($P$8),2)'
    ws['H4'] = '=ROUND(N($P$7),2)'
    ws.row_dimensions[9].height = ws.row_dimensions[8].height
    ws['A3'] = ('★ 只填 B2 货主/客户与起止日期。这里只列公司（果然鲜）真买下来的那几笔：《01》原料出库明细里出库原因＝「公司购买」'
                '并填了采购单价的行（按筐），以及成品出库明细里填了采购单价的行（按箱）；没填单价、金额为 0 的不在这里出现。'
                '「采购金额（应付）」只算货款，运费单列，两者相加＝「应付合计」。右侧数量按筐、箱分开，不能相加。')


# ════════════════════════════════════════════════════════════════════════════
# ⑤ 物料与筐子销售对账单：排除包装物料「采购入库」「期初库存」
# ════════════════════════════════════════════════════════════════════════════
_PK = re.compile(r"((?:对接源_0[12][^!]*)!\$D\$4:\$D\$(\d+)),\"包装物料\*\"")


def material_sales(wb):
    ws = wb['物料与筐子销售对账单']
    n = 0
    for r in range(3, 8):
        x = ws[f'N{r}']
        if isinstance(x.value, str) and '"包装物料*"' in x.value:
            x.value = _PK.sub(lambda m: f'{m.group(1)},"包装物料*",{m.group(1)},"<>包装物料－采购入库",'
                                        f'{m.group(1)},"<>包装物料－期初库存"', x.value)
            n += 1
    assert n == 4, n
    ws['A3'] = ('★ 只填 B2 客户与起止日期。这里是「要收钱」的那部分：包装物料的销售/领用/退回（不含采购入库、期初库存）'
                '＋ 卖断给客户的筐子，两类按业务顺序混排在同一张明细里。筐子行的「规格 / 业务类别」显示「周转筐－销售出库」。'
                '单价＝金额÷数量，自动算。★ 客户借走要还的周转筐不在这里，去《筐子对账单》看；向供应商进货看《03》往来对账单。')


# ════════════════════════════════════════════════════════════════════════════
# ⑥ 应收应付汇总
# ════════════════════════════════════════════════════════════════════════════
def ar_ap(wb):
    ws = wb['应收应付汇总']
    T0, T1, TOT = 4, 253, 254
    st_hdr = copy.copy(ws['C3']._style)
    st_num = copy.copy(ws['C4']._style)
    st_dir = copy.copy(ws['K4']._style)
    st_rem = copy.copy(ws['L4']._style)
    st_tot = copy.copy(ws['C254']._style)
    st_totl = copy.copy(ws['B254']._style)
    st_tote = copy.copy(ws['K254']._style) if ws['K254'].has_style else st_tot
    for r in range(3, TOT + 1):
        for c in range(CI('C'), CI('N') + 1):
            drop(ws, r, c)
    for m in list(ws.merged_cells.ranges):
        if m.min_row <= 2:
            ws.unmerge_cells(str(m))
    ws.merge_cells('A1:N1')
    ws.merge_cells('A2:N2')
    # G：《02》对账明细接口 K 列次果筐/托盘段只取使用费（押金收退在《02》【押金汇总】看），表头写明
    hdr = {'C': '水果成品出库\n（应收·按购买方）', 'D': '其他代存', 'E': '包装物料\n（销售/领用/退回）', 'F': '周转筐\n（销售等）',
           'G': '次果筐/托盘\n（使用费，\n不含押金）', 'H':'公司购买果品\n（应付·负数）', 'I': '物料/筐子采购\n（应付·负数）', 'J': '业务金额合计',
           'K': '已收/已付\n（付出为负）', 'L': '未收/未付', 'M': '欠款方向', 'N': '备注'}
    for c, t in hdr.items():
        put(ws, f'{c}3', t, st_hdr)
    ws.row_dimensions[3].height = 48                      # G3 三行（10 号粗体），其余两行
    an = f'{AUTO}!$AN$4:$AN${E1}'
    k1, c1, d1, l1 = R1('K'), R1('C'), R1('D'), R1('L')
    k2, c2, d2, l2 = R2('K'), R2('C'), R2('D'), R2('L')
    pk_ex = f'{d2},"<>包装物料－采购入库",{d2},"<>包装物料－期初库存"'
    zz_ex = f'{d2},"<>周转筐－采购入库",{d2},"<>周转筐－期初库存"'
    for r in range(T0, T1 + 1):
        b = f'$B{r}'
        f = {}
        f['C'] = f'SUMIFS({k1},{an},{b},{d1},"成品出库*")'
        f['D'] = f'SUMIFS({k1},{c1},{b},{d1},"其他代存*")'
        f['E'] = f'SUMIFS({k2},{c2},{b},{d2},"包装物料*",{pk_ex})'
        f['F'] = f'SUMIFS({k2},{c2},{b},{d2},"周转筐*",{zz_ex})'
        f['G'] = f'SUMIFS({k2},{c2},{b},{d2},"次果筐*")'
        f['H'] = f'-SUMIFS({R1("AJ")},{R1("Y")},{b})'
        f['I'] = f'-SUMIFS({k2},{c2},{b},{d2},"包装物料－采购入库")-SUMIFS({k2},{c2},{b},{d2},"周转筐－采购入库")'
        got = []
        for st in ('已收讫', '已付款'):
            got.append(f'SUMIFS({k1},{an},{b},{d1},"成品出库*",{l1},"{st}")')
            got.append(f'SUMIFS({k1},{c1},{b},{d1},"其他代存*",{l1},"{st}")')
            got.append(f'SUMIFS({k2},{c2},{b},{d2},"包装物料*",{pk_ex},{l2},"{st}")')
            got.append(f'SUMIFS({k2},{c2},{b},{d2},"周转筐*",{zz_ex},{l2},"{st}")')
            got.append(f'SUMIFS({k2},{c2},{b},{d2},"次果筐*",{l2},"{st}")')
        paid = (f'SUMIFS({R1("AJ")},{R1("Y")},{b},{R1("AB")},"已付清")'
                f'+SUMIFS({k2},{c2},{b},{d2},"包装物料－采购入库",{l2},"已付款")'
                f'+SUMIFS({k2},{c2},{b},{d2},"周转筐－采购入库",{l2},"已付款")')
        for c in 'CDEFGHI':
            put(ws, f'{c}{r}', f'=IF({b}="","",ROUND({f[c]},2))', st_num)
        put(ws, f'J{r}', f'=IF({b}="","",ROUND(SUM($C{r}:$I{r}),2))', st_num)
        put(ws, f'K{r}', f'=IF({b}="","",ROUND({"+".join(got)}-({paid}),2))', st_num)
        put(ws, f'L{r}', f'=IF({b}="","",ROUND(N($J{r})-N($K{r}),2))', st_num)
        put(ws, f'M{r}', f'=IF({b}="","",IF(ROUND(N($L{r}),2)>0,"对方欠款",IF(ROUND(N($L{r}),2)<0,"我方欠款","已两清")))',
            st_dir)
        put(ws, f'N{r}', None, st_rem)
        ws[f'B{r}'] = f'=IFERROR(INDEX({AUTO}!$AX$4:$AX$303,ROW()-3),"")'
    for c in 'CDEFGHIJKL':
        put(ws, f'{c}{TOT}', f'=ROUND(SUM({c}{T0}:{c}{T1}),2)', st_tot)
    put(ws, f'M{TOT}', None, st_tote)
    put(ws, f'N{TOT}', None, st_tote)
    cnt = f'N({AUTO}!$AX$1)'
    ws['A2'] = (f'=IF({cnt}>250,"★★ 注意：往来单位已有 "&{cnt}&" 个，超过本表 250 个的显示上限，请联系维护人员扩表",'
                f'"★ 全自动，全期累计（要看某段时间请用各对账单）。成品出库应收记在购买方头上（没填购买方的记货主）；'
                f'公司购买果品、包装物料/周转筐采购入库是我方应付，记负数；已收/已付里我方付出去的也记负数。'
                f'未收/未付＞0＝对方欠我方。【往来单位 "&{cnt}&"/250】'
                f'　★ 已收/已付取自《01》《02》的收付款状态，不是实际收付款；欠多少以《03》往来对账单明细为准")')
    ws.row_dimensions[2].height = 48                      # 9 号字、A2:N2 合并约 3 行
    a2 = copy.copy(ws['A2'].alignment)
    a2.wrap_text = True
    ws['A2'].alignment = a2
    cf = ws.conditional_formatting
    for rng in list(cf._cf_rules):
        for rule in cf._cf_rules[rng]:
            if rule.formula:
                rule.formula = [f'({cnt}>250)' if '_自动清单' in x else x for x in rule.formula]
    # 底稿列组是 C:J、M:XFD，直接按单列设宽会和组重叠；set_col_widths 先拆组（样式照旧，O:XFD 保持原宽 9）
    set_col_widths(ws, {'C': 14, 'D': 10, 'E': 13, 'F': 11, 'G': 11, 'H': 14, 'I': 14, 'J': 13, 'K': 13,
                        'L': 13, 'M': 10, 'N': 14})
    ws.print_area = 'A1:N18'


# ════════════════════════════════════════════════════════════════════════════
# ⑦ 主页 / 使用说明
# ════════════════════════════════════════════════════════════════════════════
def home(wb):
    ws = wb['主页']
    desc = {
        '公司购买果品对账单': '★公司（果然鲜）买下客户果品的逐笔明细与应付金额（含《01》原料出库「公司购买」，数量按筐、箱分开）',
        '物料与筐子销售对账单': '包装物料销售/领用 ＋ 卖断筐子销售，含已收/未收（不含采购入库）',
        '应收应付汇总': '各往来单位的应收、应付与净往来（成品应收按购买方；公司购买、物料筐子采购记应付）',
        '筐子库存汇总': '公司自有各类型筐子的期初、采购、退回、发出、销售与在库结存（取《02》第 5 行起）',
    }
    for r in range(6, 30):
        b = ws[f'B{r}'].value
        if isinstance(b, str) and b in desc:
            ws[f'C{r}'] = desc[b]
    ws['A2'] = ('所有对账单集中在这一本；数据全部来自《01》《02》，三个文件必须放同一个文件夹　｜　点蓝色表名可跳转　｜　'
                '可查 01 业务 6,000 笔、02 业务 3,000 笔')


def manual(wb):
    ws = wb['使用说明']
    hdr, item = ws['A48'], ws['A49']
    r0 = 52
    lines = [
        '★ 本 次 更 新（2026-10-06）',
        '1、容量：能接《01》6,000 笔、《02》3,000 笔业务（原来各 1,200 笔）。各对账单明细段放大：筐子对账单 出库/退回/随筐入库 '
        '各 300 行，自备筐对账单 入库/出库 各 300 行，水果对账单 原料入库 400、原料出库 200、成品出库 300 行，物料与筐子销售、'
        '果然鲜销售各 400 行，公司购买 300 行。',
        '2、每段标题后面显示「共 N 笔」；超过显示行数时标题变红，提示「下面只列前 M 笔」，请缩短起止日期分几次打印。'
        '黄条和右侧的合计都是按全部数据算的，不受显示行数限制；水果对账单的等级汇总也改成按全部数据算。',
        '3、【公司购买果品对账单】也列出《01》原料出库里「公司购买」并填了采购单价的行；右侧数量按筐、箱分开。',
        '4、【物料与筐子销售对账单】不再把向新顺发泡网等供应商进的货（采购入库）和期初库存当成要收的钱。',
        '5、【应收应付汇总】成品应收记在购买方头上；新增「公司购买果品（应付）」「物料/筐子采购（应付）」两列（负数），'
        '已收/已付里我方付出去的记负数。',
        '6、【筐子库存汇总】跟着《02》改了：那边第 4 行是合计，这边从第 5 行开始取。',
    ]
    for i, t in enumerate(lines):
        rr = r0 + i
        ws.merge_cells(f'A{rr}:H{rr}')
        ws[f'A{rr}'] = t
        ws[f'A{rr}']._style = copy.copy((hdr if i == 0 else item)._style)
        ws.row_dimensions[rr].height = 19.5 if i == 0 else 45


# ════════════════════════════════════════════════════════════════════════════
def finish(wb):
    """兜底：全册剩下的 $1203 / $2403 / 1200 分段一律换成新区间；下拉清单放到 300 个"""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('='):
                    nv = retarget(c.value)
                    if nv != c.value:
                        c.value = nv
        for dv in ws.data_validations.dataValidation:
            if dv.formula1 and dv.formula1.strip('=') == '_自动清单!$E$4:$E$203':
                dv.formula1 = '_自动清单!$E$4:$E$303'


def apply(wb):
    src_sheet(wb, S1, E1, N1)
    src_sheet(wb, S2, E2, N2)
    auto_list(wb)
    news = {}
    for name, spec in LAYOUT.items():
        news[name] = relayout(wb, name, spec)
    fruit_grades(wb, news['水果对账单'])
    gx_sales(wb)
    basket_stock(wb)
    company_buy(wb)
    material_sales(wb)
    ar_ap(wb)
    home(wb)
    manual(wb)
    finish(wb)


# ════════════════════════════════════════════════════════════════════════════
# 存盘后：共享公式 ＋ 中文字符引用还原
# ════════════════════════════════════════════════════════════════════════════
SHARE_SHEETS = {AUTO, '自备筐对账单', '筐子对账单', '物料与筐子销售对账单', '水果对账单', '果然鲜销售对账单',
                '公司购买果品对账单', '应收应付汇总'}


def post(path):
    """① 工作表 XML 里 &#21407; 这种字符引用的中文改回原字；
       ② 指定几张表里同一列上下行写法一样的公式改成「共享公式」（只在第一格写全文），算出来完全一样；
       ③ 核对跨文件链接关系还在（2 条）。"""
    z = zipfile.ZipFile(path)
    wbx = z.read('xl/workbook.xml').decode('utf8')
    rels = z.read('xl/_rels/workbook.xml.rels').decode('utf8')
    rid = {}
    for tag in re.findall(r'<Relationship\b[^>]*>', rels):
        i, t = re.search(r'\bId="([^"]+)"', tag), re.search(r'\bTarget="([^"]+)"', tag)
        if i and t:
            rid[i.group(1)] = t.group(1)
    part = {}
    for m in re.finditer(r'<sheet\b[^>]*>', wbx):
        nm = re.search(r'\bname="([^"]+)"', m.group(0)).group(1)
        t = rid[re.search(r'\br:id="([^"]+)"', m.group(0)).group(1)].lstrip('/')
        part[t if t.startswith('xl/') else 'xl/' + t] = nm
    n_links = rels.count('relationships/externalLink"')
    tmp = path + '.tmp'
    zo = zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED)
    for it in z.infolist():
        data = z.read(it.filename)
        if it.filename in part:
            s = _utf8_refs(data.decode('utf8'))
            if part[it.filename] in SHARE_SHEETS:
                s = _share(s)
            data = s.encode('utf8')
        zo.writestr(it, data)
    zo.close()
    z.close()
    shutil.move(tmp, path)
    if n_links != 2:
        raise SystemExit(f'《04》跨文件链接关系不是 2 条：{n_links}')
    return n_links


def _utf8_refs(s):
    def one(m):
        n = int(m.group(1)) if m.group(1) else int(m.group(2), 16)
        return chr(n) if n > 127 else m.group(0)
    return re.sub(r'&#(?:(\d+)|x([0-9a-fA-F]+));', one, s)


_FCELL = re.compile(r'<c r="([A-Z]{1,3})(\d+)"([^>]*)><f>([^<]*)</f>')
_REFX = re.compile(r'(?<![A-Za-z0-9_$.\]])(\$?)([A-Z]{1,3})(\$?)(\d+)(?![\d(A-Za-z_!])')
_STRX = re.compile(r'&quot;(?:(?!&quot;).)*&quot;|"(?:[^"]|"")*"')


def _norm(f, row):
    """公式里相对行号换成相对本行的偏移；同一列两格 _norm 相同＝可以共享"""
    out, i = [], 0
    for m in _STRX.finditer(f):
        out.append(_REFX.sub(lambda x: x.group(0) if x.group(3) else
                             f'{x.group(1)}{x.group(2)}R[{int(x.group(4)) - row}]', f[i:m.start()]))
        out.append(m.group(0))
        i = m.end()
    out.append(_REFX.sub(lambda x: x.group(0) if x.group(3) else
                         f'{x.group(1)}{x.group(2)}R[{int(x.group(4)) - row}]', f[i:]))
    return ''.join(out)


def _share(s):
    if 't="shared"' in s:
        return s
    groups, done = {}, []
    for m in _FCELL.finditer(s):
        col, row, f = m.group(1), int(m.group(2)), m.group(4)
        nf = _norm(f, row)
        g = groups.get(col)
        if g and g[1] == row - 1 and g[2] == nf:
            g[1] = row
        else:
            if g:
                done.append((col, g[0], g[1]))
            groups[col] = [row, row, nf]
    done += [(c, g[0], g[1]) for c, g in groups.items()]
    role, si = {}, 0
    for col, r0, r1 in sorted(done, key=lambda x: (CI(x[0]), x[1])):
        if r1 - r0 < 1:
            continue
        role[(col, r0)] = ('M', si, f'{col}{r0}:{col}{r1}')
        for r in range(r0 + 1, r1 + 1):
            role[(col, r)] = ('S', si, None)
        si += 1

    def rep(m):
        col, row = m.group(1), int(m.group(2))
        x = role.get((col, row))
        if not x:
            return m.group(0)
        head = f'<c r="{col}{row}"{m.group(3)}>'
        if x[0] == 'M':
            return f'{head}<f t="shared" ref="{x[2]}" si="{x[1]}">{m.group(4)}</f>'
        return f'{head}<f t="shared" si="{x[1]}"/>'
    return _FCELL.sub(rep, s)
