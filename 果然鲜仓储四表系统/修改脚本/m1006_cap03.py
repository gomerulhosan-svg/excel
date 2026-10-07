# -*- coding: utf-8 -*-
"""A050 容量（10/07）《03 财务账套与报表》：按一年用量预留行数（标签 cap03）。

用法：apply(wb) 就地修改 openpyxl 工作簿。集成时跑在 m1006_03a、m1006_03b 之后、m1006_99.apply_03 之前：
      [(m1006_03a,'apply'), (m1006_03b,'apply'), (m1006_cap03,'apply'), (m1006_99,'apply_03')]
      本模块没有 post()。对同一份底稿重复跑，结果一样（不读时间、不用随机数）。

用户原话：「A050 科目表里费用项目我增加到 124 行添加的就无效需要调到 200 行，还有资金日记账的计算行数就到 603，
          这些都要预留到一年用量，包括其他的地方检查一下也一样。」

改了什么（原容量 → 新容量）：
 1. 【科目表】数据区统一 4～203 行：全册引用 科目表!$x$4:$x$123（和 $140 / $81 / $143）的公式、下拉、条件格式、
    定义名称（报表项目表、科目名称表）、筛选区一律放到 $203（约 6 万处：资金日记账 K/M、记账凭证 F/G、期初余额、
    科目余额表、结转损益、成本费用登记表/收入成本映射/期初余额/科目余额表/记账凭证 的下拉）。
    第 124～203 行补上细灰框线的录入格样式（U 列费用对应科目补浅蓝录入色）；U、T 两列加下拉（科目名称表，
    选不到也能手打，只提醒）。下拉用的命名区（费用项目表、往来单位表…）原来按 COUNTA 算高度——中间空一行，
    最后一个就掉出下拉；改成按「最后一个非空行」算（隐藏的第 204 行是各列末行号），中间空行也不丢。
 2. 【资金日记账】清掉遗留的筛选条件（资金账户＝杨萌微信，第 4～603 行藏了 587 行，看着像只算到 603），
    筛选按钮保留 A3:V4003，第 4～4003 行全部取消隐藏，表上的「筛选中」标记（filterMode）一并清掉。
    【_自动清单】E 列往来单位池 250 个，对账补充行下拉、库存价值与欠款比对 原来只读前 200 个（$E$203），统一读到 $E$253；
    T～X（费用类别池，只读日记账 4～603、成本费用登记表）和 Z～AD（应付对象池）全册核实没人引用
    （公式、下拉、条件格式、定义名称、其他三本都查了），公式删掉——留着只会让人以为它管着 603 行的上限。
 3. 按一年用量扩（数据区 原 → 新）：
    科目余额表、期初余额  4～81（78 个）→ 4～153（150 个），合计挪到第 154 行（科目余额表 J85 自查格挪到 J157），
        资产负债表 / 结转损益 / 资金与往来余额表 / 科目余额表 引用它们的范围跟着到 $153；用户手打的科目名、金额不动；
    结转损益  Q 列（按科目表逐行数损益科目）4～123 → 4～203；
    收入成本映射  4～34（31 项）→ 4～83（80 项），下拉、进销存对接 / 结转损益 的查找范围跟着到 $83；
    费用月度支出汇总  费用项目 80 → 200（7～206，合计 207）；
    费用月度明细汇总  左 费用项目 80 → 200（7～206，合计 207）；右 逐笔 600 → 2,000（7～2006）；
    家用月度支出汇总  单位 20 → 50（9～58，合计 59）；项目 60 → 200（64～263，合计 264）；
    家用月度明细汇总  单位 20 → 50（7～56，合计 57）；项目 60 → 200（60～259，合计 260）；逐笔 600 → 2,000（7～2006）；
    应付账款明细  100 → 200 家（7～206，合计 207），辅助列 L～N 200 → 300 个往来单位；
    资金与往来余额表 往来段 150 → 300（33～332，合计 333）；应收应付汇总表、往来对账单明细 200 → 300（7～306，合计 307），
        往来对账单 / 应收 / 应付账款对账单 的单位下拉 INDIRECT("往来对账单明细!$B$7:$B$306") 跟着改；
    收入月度汇总、支出月度汇总  单位 150 → 300（7～306，合计 307）；
    往来业务明细  「02装筐费」段 500 → 1,500 行（读《02》装筐费接口 4～1503；新增 1,000 行接在 17604～18603）；
                  「02使用费」段补读《02》次果筐+托盘明细第 204～603 行（400 行，接在 18604～19003）；
                  全册 往来业务明细!…$17603 → $19003，筛选 A3:P19003。
    逐笔明细（费用/家用）的排序取数加护栏：先数「区间里有几笔」，超出的行不再跑 SMALL。
    够用、没扩的（容量 / 现用量，2026-10-05）：国外费用汇总 单位 30/1、逐笔 600/1；借款汇总 借入 30/4、借出 30/6、
        逐笔 300/16；资金与往来余额表 资金账户 20/11；_自动清单 往来单位池 250/47；库存价值与欠款比对 上半 200/47；
        成本费用登记表 300/0；对账补充行 300/1；记账凭证手工区 600/0；进销存对接 302/24；原料入库计价 300 组/26；
        结转损益 损益科目位 30/20。
 4. 【主页】最后加「容量一览」：每张录入表 / 列表 已用（公式数）/ 容量 / 占用，占用≥80% 标红；写明
    科目表费用项目可以加到第 203 行、资金日记账到第 4003 行、明细表新数据接在最后一行下面录、不要中间插行。
    【使用说明】末尾补一段。

集成约定：《03》读《02》装筐费接口 $4:$1503、次果筐+托盘明细 $4:$603（《02》那边扩到这么多行就能全接上，没扩也只是取到空）；
        往来业务明细新末行 19003（K9 的 17603 作废）；《01》《02》其他接口（成品出库明细 4～3003、公司购买接口 4～403、
        周转筐 4～3003、包装物料 4～2003）的往来段本模块没动，那几张表若再扩，往来业务明细要接着 19003 往后加段。

做法：「插入行」都按 Excel 插入单元格（下移）的规矩做——插在表的最后一个数据行上，全册指向这张表的公式 / 下拉 /
条件格式 / 定义名称 / 筛选区里，落在插入点以下的行号整体下移，跨过插入点的区域自动拉长；被挪下去的块里写死的
ROW()-k 跟着改成 ROW()-(k+n)；新行按原末行的公式（相对行号逐行平移）、样式、行高铺满。
"""
import copy
import re

from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string as CI, get_column_letter as L
from openpyxl.worksheet.cell_range import CellRange, MultiCellRange
from openpyxl.worksheet.datavalidation import DataValidation

MAXC = 16384
WARN = []                      # 插入时碰到「区域只覆盖了一部分插入列」等说不清的引用，记下来（apply 末尾要求为空）

# ════════════════════════════════════════════════════════════════════════════
#  零、公式里的引用：解析 / 改写
# ════════════════════════════════════════════════════════════════════════════
_TOK = re.compile(
    r'"(?:[^"]|"")*"'                                                   # 字符串常量：原样跳过
    r"|(?<![A-Za-z0-9_.$!:一-鿿\]'])"
    r"(?P<pfx>(?:'(?:[^']|'')+'|[A-Za-z0-9_.\[\]一-鿿]+)!)?"    # 表名前缀（含 [1] 外链）
    r"(?:(?P<c1>\$?[A-Z]{1,3})(?P<r1>\$?\d+)(?::(?P<c2>\$?[A-Z]{1,3})(?P<r2>\$?\d+))?"
    r"|(?P<q1>\$?\d+):(?P<q2>\$?\d+))"                                  # 整行区域 $1:$603
    r"(?![A-Za-z0-9_(.])")


class Ref:
    __slots__ = ('kind', 'c1', 'r1', 'c2', 'r2', 'ac1', 'ar1', 'ac2', 'ar2')

    @classmethod
    def of(cls, m):
        x = cls()
        if m.group('q1'):
            x.kind = 'rows'
            x.c1 = x.c2 = None
            x.ac1 = x.ac2 = False
            x.ar1, x.r1 = m.group('q1').startswith('$'), int(m.group('q1').lstrip('$'))
            x.ar2, x.r2 = m.group('q2').startswith('$'), int(m.group('q2').lstrip('$'))
            return x
        x.kind = 'range' if m.group('c2') else 'cell'
        c1, r1 = m.group('c1'), m.group('r1')
        x.ac1, x.c1 = c1.startswith('$'), CI(c1.lstrip('$'))
        x.ar1, x.r1 = r1.startswith('$'), int(r1.lstrip('$'))
        if x.kind == 'range':
            c2, r2 = m.group('c2'), m.group('r2')
            x.ac2, x.c2 = c2.startswith('$'), CI(c2.lstrip('$'))
            x.ar2, x.r2 = r2.startswith('$'), int(r2.lstrip('$'))
        else:
            x.c2, x.r2, x.ac2, x.ar2 = x.c1, x.r1, x.ac1, x.ar1
        return x

    def text(self):
        d = lambda a: '$' if a else ''
        if self.kind == 'rows':
            return f'{d(self.ar1)}{self.r1}:{d(self.ar2)}{self.r2}'
        a = f'{d(self.ac1)}{L(self.c1)}{d(self.ar1)}{self.r1}'
        if self.kind == 'cell':
            return a
        return a + f':{d(self.ac2)}{L(self.c2)}{d(self.ar2)}{self.r2}'


def rewrite(f, host, fn):
    """对公式 f（可带也可不带开头的 =）里每个引用调用 fn(目标表名, Ref)；fn 改了 Ref 就返回 True。
       host＝公式所在的表（没有表名前缀的引用算它的）；定义名称传 None。外链引用的表名形如 [2]装筐费接口。"""
    if not isinstance(f, str):
        return f
    out, last, hit = [], 0, False
    for m in _TOK.finditer(f):
        if m.group(0).startswith('"'):
            continue
        pfx = m.group('pfx')
        if pfx:
            nm = pfx[:-1]
            if nm.startswith("'"):
                nm = nm[1:-1].replace("''", "'")
        else:
            nm = host
        if nm is None:
            continue
        ref = Ref.of(m)
        if fn(nm, ref):
            out.append(f[last:m.start()])
            out.append((pfx or '') + ref.text())
            last = m.end()
            hit = True
    if not hit:
        return f
    out.append(f[last:])
    return ''.join(out)


def shift_rel(f, d):
    """复制公式的规矩：所有相对行号（不带 $ 的）平移 d 行（含别的表、外链）"""
    if not d:
        return f

    def fn(_nm, ref):
        ch = False
        if not ref.ar1:
            ref.r1 += d
            ch = True
        if ref.kind == 'cell':
            ref.r2 = ref.r1
        elif not ref.ar2:
            ref.r2 += d
            ch = True
        return ch
    return rewrite(f, '\0', fn)          # host 用占位名：没有表名前缀的引用也要平移


def _formula_cells(ws):
    for cell in ws._cells.values():
        v = cell.value
        if isinstance(v, str) and v.startswith('='):
            yield cell


def rewrite_wb(wb, target, fn, self_too=True, must=None):
    """全册：单元格公式、下拉、条件格式、定义名称里指向 target 表的引用，按 fn 改（must：只看含这段文字的公式）"""
    n = 0
    names = (target + '!', "'" + target + "'!")
    for ws in wb.worksheets:
        own = ws.title == target and self_too
        for cell in _formula_cells(ws):
            v = cell.value
            if must and must not in v:
                continue
            if own or any(x in v for x in names):
                nv = rewrite(v, ws.title, fn)
                if nv != v:
                    cell.value = nv
                    n += 1
        for dv in ws.data_validations.dataValidation:
            for attr in ('formula1', 'formula2'):
                v = getattr(dv, attr)
                if v and (own or any(x in v for x in names)):
                    nv = rewrite(v, ws.title, fn)
                    if nv != v:
                        setattr(dv, attr, nv)
                        n += 1
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                fl = list(rule.formula or [])
                nf = [rewrite(x, ws.title, fn) if (own or any(y in x for y in names)) else x for x in fl]
                if nf != fl:
                    rule.formula = nf
                    n += 1
    for _nm, dn in list(wb.defined_names.items()):
        v = dn.attr_text or ''
        if any(x in v for x in names):
            nv = rewrite(v, None, fn)
            if nv != v:
                dn.attr_text = nv
                n += 1
    for ws in wb.worksheets:
        for _nm, dn in list(ws.defined_names.items()):
            v = dn.attr_text or ''
            if any(x in v for x in names):
                nv = rewrite(v, None, fn)
                if nv != v:
                    dn.attr_text = nv
                    n += 1
    return n


def _range_str(s, fn, host):
    """sqref / 筛选区这类「A1:B9 C3」字符串：按 fn 改（目标表＝host）"""
    parts = []
    for p in str(s).split():
        parts.append(rewrite(p, host, fn))
    return ' '.join(parts)


# ════════════════════════════════════════════════════════════════════════════
#  一、插入行（Excel「插入单元格，下移」的规矩）与按模板铺行
# ════════════════════════════════════════════════════════════════════════════
def _ins_fn(T, at, n, clo, chi, where):
    full = clo == 1 and chi >= MAXC

    def fn(nm, ref):
        if nm != T:
            return False
        if ref.kind == 'rows':
            if full:
                ch = False
                if ref.r1 >= at:
                    ref.r1 += n
                    ch = True
                if ref.r2 >= at:
                    ref.r2 += n
                    ch = True
                return ch
            if max(ref.r1, ref.r2) >= at:
                WARN.append(f'{where}：整行引用 {ref.text()} 跨过只插部分列的 {T}')
            return False
        lo, hi = min(ref.c1, ref.c2), max(ref.c1, ref.c2)
        if clo <= lo and hi <= chi:
            ch = False
            if ref.r1 >= at:
                ref.r1 += n
                ch = True
            if ref.kind == 'cell':
                ref.r2 = ref.r1
                return ch
            if ref.r2 >= at:
                ref.r2 += n
                ch = True
            return ch
        if hi < clo or lo > chi:
            return False
        if max(ref.r1, ref.r2) >= at:
            WARN.append(f'{where}：{T}!{ref.text()} 只覆盖了插入列 {L(clo)}:{L(min(chi, MAXC))} 的一部分')
        return False
    return fn


_ROWK = re.compile(r'ROW\(\)-(\d+)')


def insert_rows(wb, sheet, at, n, clo=1, chi=MAXC, heights=True):
    """在 sheet 第 at 行（含）以下、列 clo～chi 范围内插入 n 行（单元格下移）。"""
    ws = wb[sheet]
    full = clo == 1 and chi >= MAXC
    where = f'{sheet} 插 {n} 行@{at}'
    # ① 合并区：插入列范围内、在插入点以下（或跨过插入点）的先拆开，挪完再合
    remerge = []
    for mr in list(ws.merged_cells.ranges):
        if mr.max_row < at:
            continue
        if clo <= mr.min_col and mr.max_col <= chi:
            remerge.append((mr.min_col, mr.min_row, mr.max_col, mr.max_row))
            ws.unmerge_cells(str(mr))
        elif mr.max_col < clo or mr.min_col > chi:
            continue
        else:
            raise SystemExit(f'{where}：合并区 {mr} 横跨插入列边界')
    # ② 单元格从下往上挪；被挪到下面的块（原行号 > at）里写死的 ROW()-k 改成 ROW()-(k+n)
    keys = sorted((k for k in ws._cells if k[0] >= at and clo <= k[1] <= chi), reverse=True)
    for (r, c) in keys:
        cell = ws._cells.pop((r, c))
        cell.row = r + n
        ws._cells[(r + n, c)] = cell
        v = cell.value
        if r > at and isinstance(v, str) and v.startswith('=') and 'ROW()' in v:
            nv = _ROWK.sub(lambda m: f'ROW()-{int(m.group(1)) + n}', v)
            if re.search(r'ROW\(\)(?!-\d)', nv):
                WARN.append(f'{where}：{cell.coordinate} 有不是 ROW()-k 形式的 ROW()：{nv[:80]}')
            cell.value = nv
    # ③ 全册引用
    fn = _ins_fn(sheet, at, n, clo, chi, where)
    rewrite_wb(wb, sheet, fn)
    # ④ 本表的下拉范围、条件格式范围、筛选区、打印区
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(_range_str(dv.sqref, fn, sheet))
    if len(ws.conditional_formatting):
        old = ws.conditional_formatting
        new = ConditionalFormattingList()
        for cf in old:
            rng = _range_str(cf.sqref, fn, sheet)
            for rule in cf.rules:
                new.add(rng, rule)
        ws.conditional_formatting = new
    if ws.auto_filter.ref:
        ws.auto_filter.ref = _range_str(ws.auto_filter.ref, fn, sheet)
    if ws.print_area:
        ws.print_area = [_range_str(str(a).split('!')[-1].replace('$', ''), fn, sheet) for a in ws.print_area]
    # ⑤ 合并区按新位置合回去
    for (c1, r1, c2, r2) in remerge:
        r1n = r1 + n if r1 >= at else r1
        r2n = r2 + n if r2 >= at else r2
        ws.merge_cells(start_row=r1n, start_column=c1, end_row=r2n, end_column=c2)
    # ⑥ 行高（整行插入、或者插入列以外都是隐藏辅助列时）跟着挪
    if heights:
        rds = ws.row_dimensions
        for r in sorted((k for k in list(rds.keys()) if k >= at), reverse=True):
            rd = rds.pop(r)
            rd.index = r + n
            rds[r + n] = rd


def fill_rows(ws, cols, tmpl, rows, values=True, height=True, post=None):
    """按第 tmpl 行铺 rows：公式按复制规矩平移相对行号，常量照抄（values=False 时只抄公式），样式照抄。
       post(列号, 行号, 公式) 可以再改写（处理写死的常数）。"""
    src = {c: ws.cell(tmpl, c) for c in cols}
    h = ws.row_dimensions[tmpl].height
    for r in rows:
        d = r - tmpl
        for c, s in src.items():
            dst = ws.cell(r, c)
            v = s.value
            if isinstance(v, str) and v.startswith('='):
                v = shift_rel(v, d)
                if post:
                    v = post(c, r, v)
                dst.value = v
            elif values:
                dst.value = v
            else:
                dst.value = None
            if s.has_style:
                dst._style = copy.copy(s._style)
        if height and h:
            ws.row_dimensions[r].height = h


def _cols(a, b):
    return list(range(CI(a), CI(b) + 1))


# ════════════════════════════════════════════════════════════════════════════
#  二、科目表：全册引用统一到第 203 行；第 124～203 行补样式和下拉；命名区不怕中间空行
# ════════════════════════════════════════════════════════════════════════════
KM0, KM1 = 4, 203
KM_HELP = 204                                   # 隐藏行：各列最后一个非空行（数据区内的序号）


def kemu(wb):
    def fn(nm, ref):
        if nm != '科目表' or ref.kind != 'range':
            return False
        if min(ref.r1, ref.r2) == KM0 and max(ref.r1, ref.r2) in (81, 123, 140, 143):
            ref.r2 = KM1
            return True
        return False
    n = rewrite_wb(wb, '科目表', fn)
    ws = wb['科目表']
    ws.auto_filter.ref = f'A3:G{KM1}'
    # 样式：没框线的格子按本列第 4 行补（A～G 科目、I～Q 各码表、T 资金账户对应科目、U 费用对应科目）
    cols = _cols('A', 'G') + _cols('I', 'Q') + [CI('T'), CI('U')]
    for c in cols:
        s4 = ws.cell(KM0, c)
        for r in range(KM0, KM1 + 1):
            d = ws.cell(r, c)
            if d.border.left.style is None or (c == CI('U') and d.fill.fill_type is None):
                d._style = copy.copy(s4._style)
    # 下拉用的命名区：高度从 COUNTA（中间空一行就把最后一个挤掉）改成「最后一个非空行」
    for c in range(1, CI('Q') + 1):
        col = L(c)
        ws.cell(KM_HELP, c).value = f'=SUMPRODUCT(MAX(({col}{KM0}:{col}{KM1}<>"")*ROW({col}{KM0}:{col}{KM1})))-{KM0 - 1}'
        ws.cell(KM_HELP, c).font = Font(name='宋体', size=8, color='BFBFBF')
        ws.cell(KM_HELP, c).border = Border()
        ws.cell(KM_HELP, c).fill = PatternFill(fill_type=None)
    ws.cell(KM_HELP, CI('R')).value = '← 下拉用：各列最后一行（自动，勿删）'
    ws.cell(KM_HELP, CI('R')).font = Font(name='宋体', size=8, color='BFBFBF')
    ws.row_dimensions[KM_HELP].hidden = True
    rx = re.compile(r'COUNTA\(INDEX\(科目表!\$A\$4:\$Q\$203,0,(MATCH\("[^"]+",科目表!\$A\$3:\$Q\$3,0\))\)\)')
    k = 0
    for _nm, dn in wb.defined_names.items():
        t = dn.attr_text or ''
        t2 = rx.sub(lambda m: f'INDEX(科目表!$A${KM_HELP}:$Q${KM_HELP},{m.group(1)})', t)
        if t2 != t:
            dn.attr_text = t2
            k += 1
    assert k == 11, k
    # U 费用对应科目、T 资金账户对应科目：下拉选科目（选不到也能手打，只提醒）
    for col, what in (('U', '费用对应科目'), ('T', '资金账户对应科目')):
        dv = DataValidation(type='list', formula1='科目名称表', allow_blank=True, showErrorMessage=True,
                            errorStyle='warning', errorTitle=what,
                            error='这个名字不在左边 B 列「科目名称」里，凭证会记到一个科目余额表里没有的科目。确定要这样填吗？')
        dv.add(f'{col}{KM0}:{col}{KM1}')
        ws.add_data_validation(dv)
    a2 = ws['A2'].value or ''
    add = '各列都可以录到第 203 行（中间空了行也认）。'
    if add not in a2:
        ws['A2'] = a2 + add
        ws.row_dimensions[2].height = max(ws.row_dimensions[2].height or 0, 45)
    return n


# ════════════════════════════════════════════════════════════════════════════
#  三、资金日记账：清遗留筛选、取消隐藏；_自动清单 删无人引用的老池子
# ════════════════════════════════════════════════════════════════════════════
def journal_filter(wb):
    ws = wb['资金日记账']
    ws.auto_filter.ref = 'A3:V4003'
    ws.auto_filter.filterColumn = []
    ws.auto_filter.sortState = None
    ws.sheet_properties.filterMode = False          # WPS 记的「筛选中」标记，不清的话打开还显示漏斗
    k = 0
    for r, rd in list(ws.row_dimensions.items()):
        if rd.hidden:
            rd.hidden = False
            k += 1
    return k


LEGACY = [('T', 'X', 903), ('Z', 'AD', 303)]


def _legacy_refs(wb):
    """全册（公式、下拉、条件格式、定义名称）里有没有引用 _自动清单 T～X / Z～AD 的（那几列自己除外）"""
    hits = []
    bad = set()
    for a, b, _ in LEGACY:
        bad.update(range(CI(a), CI(b) + 1))

    def mk(where):
        def fn(nm, ref):
            if nm != '_自动清单':
                return False
            lo, hi = (ref.c1, ref.c2) if ref.kind != 'rows' else (1, MAXC)
            if any(lo <= c <= hi for c in bad):
                hits.append(f'{where}: {ref.text()}')
            return False
        return fn
    for ws in wb.worksheets:
        for cell in _formula_cells(ws):
            if ws.title == '_自动清单' and cell.column in bad:
                continue
            if ws.title == '_自动清单' or '_自动清单' in cell.value:
                rewrite(cell.value, ws.title, mk(f'{ws.title}!{cell.coordinate}'))
                if 'INDIRECT' in cell.value and '_自动清单' in cell.value:
                    hits.append(f'{ws.title}!{cell.coordinate} INDIRECT')
        for dv in ws.data_validations.dataValidation:
            for v in (dv.formula1, dv.formula2):
                if v:
                    rewrite(v, ws.title, mk(f'{ws.title}[下拉 {dv.sqref}]'))
                    if 'INDIRECT' in v and '_自动清单' in v:
                        hits.append(f'{ws.title}[下拉 {dv.sqref}] INDIRECT')
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                for v in rule.formula or []:
                    rewrite(v, ws.title, mk(f'{ws.title}[条件格式 {cf.sqref}]'))
    for nm, dn in wb.defined_names.items():
        rewrite(dn.attr_text or '', None, mk(f'[名称 {nm}]'))
        if '_自动清单' in (dn.attr_text or '') and 'INDIRECT' in (dn.attr_text or ''):
            hits.append(f'[名称 {nm}] INDIRECT')
    return hits


def pool_e(wb):
    """_自动清单 E 列往来单位池有 250 个（4～253），对账补充行下拉、库存价值与欠款比对只读到 $203：统一读到 $253"""
    def fn(nm, ref):
        if nm == '_自动清单' and ref.kind == 'range' and ref.c1 == ref.c2 == CI('E') and ref.r1 == 4 and ref.r2 == 203:
            ref.r2 = 253
            return True
        return False
    return rewrite_wb(wb, '_自动清单', fn, self_too=False)


def legacy_pools(wb):
    hits = _legacy_refs(wb)
    if hits:
        raise SystemExit('_自动清单 T～X / Z～AD 还有人引用，不能删：' + '；'.join(hits[:5]))
    ws = wb['_自动清单']
    n = 0
    for a, b, r1 in LEGACY:
        for r in range(1, r1 + 1):
            for c in range(CI(a), CI(b) + 1):
                cell = ws._cells.get((r, c))
                if cell is not None and cell.value is not None:
                    cell.value = None
                    n += 1
    a2 = ws['A2'].value or ''
    if '费用类别及' in a2:
        ws['A2'] = a2.replace('往来单位、费用类别及收入往来单位清单',
                              '往来单位及收入往来单位清单（原 T～AD 列费用类别 / 应付对象池没有任何表引用，10/07 已清掉）')
    return n


# ════════════════════════════════════════════════════════════════════════════
#  四、各表扩行
# ════════════════════════════════════════════════════════════════════════════
def expand(wb, sheet, last, new_n, cols, clo=1, chi=MAXC, heights=True, post=None):
    """表的数据区最后一行是 last，在它上面插 new_n 行，再按挪下去的原末行把空出来的行铺满。返回新末行。"""
    insert_rows(wb, sheet, last, new_n, clo, chi, heights)
    ws = wb[sheet]
    fill_rows(ws, cols, last + new_n, range(last, last + new_n), values=False, post=post)
    return last + new_n


def kemu_balance(wb):
    # 科目余额表：4～81 → 4～153，合计 82 → 154（J85 自查 → J157）
    expand(wb, '科目余额表', 81, 72, _cols('A', 'J'))
    # 期初余额：4～81 → 4～153，合计 82 → 154；右边资金账户 H:I（4～43）不动
    expand(wb, '期初余额', 81, 72, _cols('A', 'F'))
    for nm in ('科目余额表', '期初余额'):
        ws = wb[nm]
        assert ws['B154'].value == '合计', (nm, ws['B154'].value)
    assert wb['科目余额表']['J157'].value == '=J154-I154', wb['科目余额表']['J157'].value
    assert wb['期初余额']['D154'].value == '=ROUND(SUM(D4:D153),2)'


def jiezhuan_q(wb):
    ws = wb['结转损益']
    assert ws['Q123'].value.startswith('=N(Q122)+IF(INDEX(科目表!$C$4:$C$203,ROW()-3)'), ws['Q123'].value
    fill_rows(ws, [CI('Q')], 123, range(124, KM1 + 1), height=False)

    def fn(nm, ref):
        if nm != '结转损益' or ref.c1 != CI('Q') or ref.c2 != CI('Q'):
            return False
        ch = False
        if ref.ar2 and ref.r2 == 123:
            ref.r2 = KM1
            ch = True
        if ref.kind == 'cell':
            ref.r1 = ref.r2
        return ch
    for cell in _formula_cells(ws):
        if cell.column != CI('Q'):
            v = cell.value
            nv = rewrite(v, ws.title, fn)
            if nv != v:
                cell.value = nv


def mapping(wb):
    # 收入成本映射：4～34 → 4～83（没有合计行；空行只要样式）
    insert_rows(wb, '收入成本映射', 34, 49)
    ws = wb['收入成本映射']
    fill_rows(ws, _cols('A', 'G'), 83, range(34, 83), values=False)


def fee_tables(wb):
    # 费用月度支出汇总：费用项目 7～86 → 7～206，合计 87 → 207（右边 X～Z 隐藏辅助列不动）
    expand(wb, '费用月度支出汇总', 86, 120, _cols('A', 'N'), 1, CI('W'))
    # 费用月度明细汇总：左 7～86 → 7～206（只动 A～D）；右 逐笔 7～606 → 7～2006
    expand(wb, '费用月度明细汇总', 86, 120, _cols('A', 'D'), 1, CI('D'), heights=False)
    _detail(wb['费用月度明细汇总'], 'E', 'L', 'S', 'R', 606, 2006)


def _detail(ws, a, b, key, sk, last, new_last):
    """右边逐笔明细铺到 new_last 行；排序取数列（key，读排序键 sk）加护栏：先数有几笔，超出的行不跑 SMALL"""
    cnt = f'{key}5'
    assert ws[cnt].value is None and ws[f'{key}4'].value is None, ws.title
    ws[f'{key}4'] = '逐笔数'
    ws[f'{key}5'] = f'=COUNT(${sk}$4:${sk}$4003)'
    fill_rows(ws, _cols(a, b), last, range(last + 1, new_last + 1))
    fill_rows(ws, [CI(key)], last, range(last + 1, new_last + 1), height=False)
    old = (f'=IFERROR(SMALL(${sk}$4:${sk}$4003,ROW()-6)-INT(SMALL(${sk}$4:${sk}$4003,ROW()-6)/10000)*10000,"")')
    new = (f'=IF(ROW()-6>${key}$5,"",IFERROR(SMALL(${sk}$4:${sk}$4003,ROW()-6)'
           f'-INT(SMALL(${sk}$4:${sk}$4003,ROW()-6)/10000)*10000,""))')
    for r in range(7, new_last + 1):
        c = ws[f'{key}{r}']
        assert c.value == old, (ws.title, r, c.value)
        c.value = new


def home_tables(wb):
    # 家用月度支出汇总：单位 9～28 → 9～58（合计 59）；项目 34～93 → 下移 30 后 64～123 → 64～263（合计 264）
    e1 = expand(wb, '家用月度支出汇总', 28, 30, _cols('A', 'N'), 1, CI('O'))
    expand(wb, '家用月度支出汇总', 93 + 30, 140, _cols('A', 'N'), 1, CI('O'))
    ws = wb['家用月度支出汇总']
    assert e1 == 58 and ws['A59'].value.startswith('合') and ws['A264'].value.startswith('合'), ws['A264'].value
    assert ws['A64'].value.startswith('=IF(ROW()-63>$U$4003'), ws['A64'].value
    # 家用月度明细汇总：单位 7～26 → 7～56（合计 57）；项目 30～89 → 60～119 → 60～259（合计 260）；逐笔 7～606 → 7～2006
    expand(wb, '家用月度明细汇总', 26, 30, _cols('A', 'D'), 1, CI('E'))
    expand(wb, '家用月度明细汇总', 89 + 30, 140, _cols('A', 'D'), 1, CI('E'))
    ws = wb['家用月度明细汇总']
    assert ws['A57'].value.startswith('合') and ws['A59'].value == '费用项目' and ws['A260'].value.startswith('合')
    assert ws['A60'].value.startswith('=IF(ROW()-59>$V$4003'), ws['A60'].value
    _detail(ws, 'F', 'N', 'X', 'W', 606, 2006)


def ap_detail(wb):
    # 应付账款明细：显示 7～106 → 7～206（合计 107 → 207）；辅助 L～N 7～206 → 7～306（300 个往来单位）
    expand(wb, '应付账款明细', 106, 100, _cols('A', 'G'), 1, CI('J'))
    ws = wb['应付账款明细']

    def post(c, r, f):
        if c == CI('L'):
            f = re.sub(r'IF\(\d+>', f'IF({r - 6}>', f)
            f = re.sub(r'MATCH\(\d+,', f'MATCH({r - 6},', f)
        return f
    fill_rows(ws, _cols('L', 'N'), 206, range(207, 307), height=False, post=post)

    def fn(nm, ref):
        if nm != '应付账款明细' or ref.kind == 'rows':
            return False
        if not (CI('L') <= ref.c1 <= CI('N')) or not ref.ar2 or ref.r2 != 206:
            return False
        ref.r2 = 306
        if ref.kind == 'cell':
            ref.r1 = 306
        return True
    for cell in _formula_cells(ws):
        if cell.column < CI('L'):
            v = cell.value
            nv = rewrite(v, ws.title, fn)
            if nv != v:
                cell.value = nv
    assert ws['B7'].value == '=IF(ROW()-6>$N$306,"",INDEX($L$7:$L$306,MATCH(ROW()-6,$N$7:$N$306,0)))', ws['B7'].value
    assert ws['L306'].value.startswith('=IF(300>往来业务明细!$Q$'), ws['L306'].value


def wl_lists(wb):
    expand(wb, '资金与往来余额表', 182, 150, _cols('A', 'L'))
    assert wb['资金与往来余额表']['A333'].value.startswith('合')
    for nm in ('应收应付汇总表', '往来对账单明细'):
        expand(wb, nm, 206, 100, _cols('A', 'L'))
        assert wb[nm]['A307'].value.startswith('合'), nm
    for nm in ('往来对账单', '应收账款对账单', '应付账款对账单'):
        k = 0
        for dv in wb[nm].data_validations.dataValidation:
            if dv.formula1 and 'INDIRECT("往来对账单明细!$B$7:$B$206")' in dv.formula1:
                dv.formula1 = dv.formula1.replace('$B$206', '$B$306')
                k += 1
        assert k == 1, (nm, k)
    for nm in ('收入月度汇总', '支出月度汇总'):
        expand(wb, nm, 156, 150, _cols('A', 'N'), 1, CI('W'))
        ws = wb[nm]
        assert ws['A307'].value.startswith('合') and '$N$307' in ws['D4'].value, nm


# ════════════════════════════════════════════════════════════════════════════
#  五、往来业务明细：02装筐费 1,500 行、02使用费 补 400 行；全册 $17603 → $19003
# ════════════════════════════════════════════════════════════════════════════
WL_OLD, WL_END = 17603, 19003
ZF0, ZF1_OLD, ZF1 = 17104, 17603, 18603          # 02装筐费：装筐费接口 4～503 → 4～1503
SY0, SY1 = 18604, 19003                          # 02使用费 续：次果筐+托盘明细 204～603
CG = "'[2]次果筐+托盘明细'"


def wanglai(wb):
    def fn(nm, ref):
        if nm != '往来业务明细':
            return False
        ch = False
        if ref.ar2 and ref.r2 == WL_OLD:
            ref.r2 = WL_END
            ch = True
        if ref.kind == 'cell':
            ref.r1 = ref.r2
        elif ref.ar1 and ref.r1 == WL_OLD:
            ref.r1 = WL_END
            ch = True
        return ch
    n = rewrite_wb(wb, '往来业务明细', fn, must=str(WL_OLD))
    ws = wb['往来业务明细']
    assert ws.auto_filter.ref == f'A3:P{WL_OLD}', ws.auto_filter.ref
    ws.auto_filter.ref = f'A3:P{WL_END}'
    # 02装筐费 原段的取数范围 $503 → $1503
    def fz(nm, ref):
        if nm == '[2]装筐费接口' and ref.kind == 'range' and ref.r2 == 503:
            ref.r2 = 1503
            return True
        return False
    for r in range(ZF0, ZF1_OLD + 1):
        assert ws[f'A{r}'].value == '02装筐费'
        for c in range(3, 17):
            cell = ws.cell(r, c)
            if isinstance(cell.value, str) and cell.value.startswith('='):
                cell.value = rewrite(cell.value, ws.title, fz)
    assert ws['C17603'].value == '=IFERROR(1*INDEX([2]装筐费接口!$B$4:$B$1503,ROW()-17103),0)', ws['C17603'].value
    # 02装筐费 续：17604～18603（接口第 504～1503 行）
    cols = _cols('A', 'X')
    fill_rows(ws, cols, ZF1_OLD, range(ZF1_OLD + 1, ZF1 + 1))
    for r in range(ZF1_OLD + 1, ZF1 + 1):
        ws[f'B{r}'] = r - (ZF0 - 4)
    # 02使用费 续：18604～19003（次果筐+托盘明细第 204～603 行），C～P 照 4204 行的写法，改成按位置取
    t = 4204
    assert ws[f'A{t}'].value == '02使用费' and ws[f'B{t}'].value == 203
    off = SY0 - 201                                 # ROW()-off＝区域 4～603 里的第几个（18604 → 201 → 第 204 行）
    rx = re.compile(re.escape(CG) + r'!([A-Z]{1,3})203(?!\d)')
    tm = {}
    for c in range(1, 17):
        v = ws.cell(t, c).value
        if isinstance(v, str) and v.startswith('='):
            v = rx.sub(lambda m: f'INDEX({CG}!${m.group(1)}$4:${m.group(1)}$603,ROW()-{off})', v)
            assert '203' not in v.replace(f'ROW()-{off}', ''), v
        tm[c] = v
    h = ws.row_dimensions[t].height
    for r in range(SY0, SY1 + 1):
        for c in range(1, 17):
            v = tm[c]
            if isinstance(v, str) and v.startswith('='):
                v = shift_rel(v, r - t)
            d = ws.cell(r, c)
            d.value = v
            d._style = copy.copy(ws.cell(t, c)._style)
        ws[f'B{r}'] = r - (SY0 - 204)
        if h:
            ws.row_dimensions[r].height = h
    fill_rows(ws, _cols('Q', 'X'), ZF1_OLD, range(SY0, SY1 + 1), height=False)
    assert ws[f'C{SY0}'].value == f"=IFERROR(1*INDEX({CG}!$B$4:$B$603,ROW()-{off}),0)", ws[f'C{SY0}'].value
    assert ws[f'K{SY0}'].value == f'=J{SY0}' and ws[f'B{SY1}'].value == 603
    assert ws[f'Q{SY1}'].value == f'=N(Q{SY1 - 1})+IF(D{SY1}="",0,IF(MATCH(D{SY1},$D$4:$D${WL_END},0)=ROW()-3,1,0))'
    assert ws[f'C{ZF1}'].value == '=IFERROR(1*INDEX([2]装筐费接口!$B$4:$B$1503,ROW()-17103),0)' and ws[f'B{ZF1}'].value == 1503
    return n


# ════════════════════════════════════════════════════════════════════════════
#  六、主页「容量一览」、使用说明
# ════════════════════════════════════════════════════════════════════════════
def _cnt(rng):
    return f'COUNTIF({rng},"?*")+COUNT({rng})'


CAPS = [
    # (子表, 区域与说明, 已用（公式，不带 =）, 容量)
    ('资金日记账', '第 4～4003 行（按 C 列资金账户数笔数）', 'COUNTA(资金日记账!$C$4:$C$4003)', 4000),
    ('科目表', '费用项目 M 列 第 4～203 行（U 列对应科目同行）', 'COUNTA(科目表!$M$4:$M$203)', 200),
    ('科目表', '会计科目 A～G 第 4～203 行', 'COUNTA(科目表!$B$4:$B$203)', 200),
    ('科目表', '往来单位 I 列 第 4～203 行', 'COUNTA(科目表!$I$4:$I$203)', 200),
    ('期初余额', '第 4～153 行（科目名称）', 'COUNTA(期初余额!$B$4:$B$153)', 150),
    ('科目余额表', '第 4～153 行（科目名称）', 'COUNTA(科目余额表!$B$4:$B$153)', 150),
    ('收入成本映射', '第 4～83 行（项目）', 'COUNTA(收入成本映射!$B$4:$B$83)', 80),
    ('成本费用登记表', '第 4～303 行（供应商/对象）', 'COUNTA(成本费用登记表!$C$4:$C$303)', 300),
    ('对账补充行', '第 4～303 行（往来单位）', 'COUNTA(对账补充行!$C$4:$C$303)', 300),
    ('记账凭证', '手工区 第 3010～3609 行（科目名称）', 'COUNTA(记账凭证!$E$3010:$E$3609)', 600),
    ('进销存对接', '第 4～305 行（《01》《02》对接源合并）', 'COUNTIF(进销存对接!$B$4:$B$305,"?*")', 302),
    ('原料入库计价', '货主＋品种 300 组（第 5～904 行，每组 3 行）', 'ROUND(COUNTIF(原料入库计价!$B$5:$B$904,"?*")/3,0)', 300),
    ('往来对账单明细', '往来单位 300 家（应收应付汇总表、资金与往来余额表 同 300 家）', 'N(往来业务明细!$Q$19003)', 300),
    ('应付账款明细', '有应付业务的单位 200 家', 'N(应付账款明细!$N$306)', 200),
    ('收入月度汇总', '有收入的往来单位 300 家（按第 2 行选的年份）', 'N(收入月度汇总!$Z$4003)', 300),
    ('支出月度汇总', '有支出的往来单位 300 家（按第 2 行选的年份）', 'N(支出月度汇总!$Z$4003)', 300),
    ('费用月度支出汇总', '费用项目 200 个（按第 2 行选的年份）', 'N(费用月度支出汇总!$Z$4003)', 200),
    ('费用月度明细汇总', '左 费用项目 200 个（按选的区间）', 'N(费用月度明细汇总!$Q$4003)', 200),
    ('费用月度明细汇总', '右 逐笔明细 2,000 笔（按选的区间）', 'N(费用月度明细汇总!$S$5)', 2000),
    ('家用月度支出汇总', '往来单位 50 个（按选的年份）', 'N(家用月度支出汇总!$R$4003)', 50),
    ('家用月度支出汇总', '费用项目 200 个（按选的年份）', 'N(家用月度支出汇总!$U$4003)', 200),
    ('家用月度明细汇总', '左上 往来单位 50 个（按选的区间）', 'N(家用月度明细汇总!$S$4003)', 50),
    ('家用月度明细汇总', '左下 费用项目 200 个（按选的区间）', 'N(家用月度明细汇总!$V$4003)', 200),
    ('家用月度明细汇总', '右 逐笔明细 2,000 笔（按选的区间）', 'N(家用月度明细汇总!$X$5)', 2000),
    ('国外费用汇总', '往来单位 30 个（左）', 'N(国外费用汇总!$S$4003)', 30),
    ('国外费用汇总', '逐笔明细 600 笔（右）', 'COUNT(国外费用汇总!$T$4:$T$4003)', 600),
    ('借款汇总', '我方借入 出借人 30 个', 'N(借款汇总!$X$4023)', 30),
    ('借款汇总', '我方借出 借款人 30 个', 'N(借款汇总!$AA$4023)', 30),
    ('借款汇总', '逐笔明细 300 笔（按选的区间）', 'COUNT(借款汇总!$AB$4:$AB$4003)', 300),
    ('库存价值与欠款比对', '上半段 客户 200 家', 'N(_自动清单!$A$1)', 200),
    ('往来业务明细', '「02装筐费」段 1,500 笔（第 17104～18603 行 ← 《02》装筐费接口）',
     'COUNTIF(往来业务明细!$C$17104:$C$18603,">0")', 1500),
    ('往来业务明细', '「02使用费」段 600 笔（第 4005～4204、18604～19003 行 ← 《02》次果筐+托盘明细）',
     'COUNTIF(往来业务明细!$C$4005:$C$4204,">0")+COUNTIF(往来业务明细!$C$18604:$C$19003,">0")', 600),
    ('结转损益', '损益科目 30 个位置（超了多出的科目不结转）', 'N(结转损益!$Q$203)', 30),
]


def capacity_panel(wb):
    ws = wb['主页']
    r0 = ws.max_row + 2
    st_sec, st_hdr, st_row = ws['A65'], ws['A66'], {c: ws[f'{c}67'] for c in 'ABCD'}
    ws.merge_cells(f'A{r0}:D{r0}')
    ws[f'A{r0}'] = '容 量 一 览（按一年用量预留 · 占用≥80% 标红）'
    ws[f'A{r0}']._style = copy.copy(st_sec._style)
    for c in 'BCD':
        ws[f'{c}{r0}']._style = copy.copy(ws[f'{c}65']._style)
    ws.row_dimensions[r0].height = 19.5
    r = r0 + 1
    ws.merge_cells(f'A{r}:D{r}')
    ws[f'A{r}'] = ('★ 科目表的费用项目（M 列）和费用对应科目（U 列）可以一直加到第 203 行；资金日记账录到第 4003 行；'
                   '所有明细表新数据一律接在最后一行下面录，不要在中间插行（插行会让凭证、往来、汇总对不上）。'
                   '下面「已用」是公式实时数的，到 80% 变红就该找维护的人扩了。')
    ws[f'A{r}'].font = Font(name='微软雅黑', size=9, bold=True, color='C00000')
    ws[f'A{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 42
    r += 1
    for c, t in zip('ABCD', ('序号', '子表名称', '区域 ｜ 已用 / 容量', '占用')):
        ws[f'{c}{r}'] = t
        ws[f'{c}{r}']._style = copy.copy(ws[f'{c}66']._style)
    ws.row_dimensions[r].height = 19.5
    first = r + 1
    for i, (sh, txt, used, cap) in enumerate(CAPS):
        r = first + i
        ws[f'A{r}'] = i + 1
        ws[f'B{r}'] = f'=HYPERLINK("#\'{sh}\'!A1","{sh}")'
        ws[f'C{r}'] = f'="{txt} ｜ 已用 "&TEXT({used},"#,##0")&" / {cap:,}"'
        ws[f'D{r}'] = f'=ROUND(({used})/{cap},4)'
        for c in 'ABCD':
            ws[f'{c}{r}']._style = copy.copy(st_row[c]._style)
        ws[f'D{r}'].number_format = '0%'
        ws[f'C{r}'].alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        units = sum(2 if ord(ch) > 255 else 1.1 for ch in f'{txt} ｜ 已用 0,000 / {cap:,}')
        ws.row_dimensions[r].height = 18 if units <= 60 else 32
    last = r
    red = Font(name='微软雅黑', size=10, bold=True, color='C00000')
    ws.conditional_formatting.add(f'D{first}:D{last}',
                                  CellIsRule(operator='greaterThanOrEqual', formula=['0.8'],
                                             fill=PatternFill('solid', fgColor='FFC7CE'), font=red))
    # 旧说明里的末行号
    for rr in range(1, first):
        v = ws[f'C{rr}'].value
        if isinstance(v, str) and '最后一行 17603' in v:
            ws[f'C{rr}'] = v.replace('最后一行 17603', '最后一行 19003（10/07 装筐费段扩到 1,500 行、使用费段补到 600 行）')
    return first, last


def manual(wb):
    ws = wb['使用说明']
    r0 = ws.max_row + 2
    lines = [
        '★ 本 次 更 新（2026-10-07 · 容量按一年预留）',
        '1、【科目表】每一列都能录到第 203 行（原来费用项目第 124 行以后录了，日记账、凭证、科目余额表取不到对应科目）。'
        '费用项目（M 列）旁边 U 列填它记到哪个科目，空着＝管理费用；U、T 两列有下拉。中间空一行也没关系，下拉照样能选到最后一个。',
        '2、【资金日记账】一直是第 4～4003 行都在算。之前看着「只到 603 行」，是表上留着一个筛选（资金账户＝杨萌微信），'
        '把 4～603 行里其他账户的 587 行藏起来了；现在筛选条件清掉、行全部显示，筛选按钮还在。',
        '3、按一年用量扩：科目余额表/期初余额 150 个科目；收入成本映射 80 项；费用月度支出汇总、费用月度明细汇总 费用项目 200 个，'
        '逐笔 2,000 笔；家用两张 单位 50、项目 200、逐笔 2,000；收入/支出月度汇总、往来对账单明细、应收应付汇总表、'
        '资金与往来余额表 往来单位 300 家；应付账款明细 200 家；往来业务明细 装筐费 1,500 行、次果筐使用费 600 行（末行 19003）。',
        '4、【主页】最下面「容量一览」实时显示每张表用了多少，到 80% 变红。所有明细表新数据只在最后一行下面接着录，不要在中间插行。',
    ]
    width = sum((ws.column_dimensions[c].width if c in ws.column_dimensions and ws.column_dimensions[c].width else 9)
                for c in 'AB') + 6 * 22
    for i, t in enumerate(lines):
        r = r0 + i
        ws.merge_cells(f'A{r}:H{r}')
        ws[f'A{r}'] = t
        ws[f'A{r}']._style = copy.copy(ws['A133' if i == 0 else 'A134']._style)
        if i == 0:
            ws.row_dimensions[r].height = ws.row_dimensions[133].height or 20
        else:
            units = sum(2 if ord(ch) > 255 else 1.1 for ch in t)
            ws.row_dimensions[r].height = max(ws.row_dimensions[134].height or 30, (int(units // (width * 0.95)) + 1) * 15 + 4)


# ════════════════════════════════════════════════════════════════════════════
def apply(wb):
    del WARN[:]
    assert wb['往来业务明细'].auto_filter.ref == 'A3:P17603', '要在 m1006_03a 之后跑'
    assert wb['库存价值与欠款比对']['A208'].value is not None, '要在 m1006_03b 之后跑'
    kemu(wb)
    journal_filter(wb)
    legacy_pools(wb)
    pool_e(wb)
    kemu_balance(wb)
    jiezhuan_q(wb)
    mapping(wb)
    fee_tables(wb)
    home_tables(wb)
    ap_detail(wb)
    wl_lists(wb)
    wanglai(wb)
    capacity_panel(wb)
    manual(wb)
    if WARN:
        raise SystemExit('插入行时有说不清的引用：\n' + '\n'.join(WARN[:20]))
