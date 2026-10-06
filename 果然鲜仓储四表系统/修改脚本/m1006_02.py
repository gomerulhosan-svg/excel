# -*- coding: utf-8 -*-
"""A050 1006 这一轮《02 物料与周转物台账》的改动（底稿：参考/1006上传/02_物料与周转物台账模板.xlsx）。

用法：apply(wb) 就地修改 openpyxl 工作簿（openpyxl.load_workbook(底稿)，不用 data_only）；
      不需要 XML 层后处理（post 留空）。对同一份底稿重复运行，结果完全一样。

这一轮改了什么：
 1. 【客户领用物料汇总】C～T 物料列改成「按本列第 3 行表头文字求和」，表头变蓝色可改（默认 香梨网、苹果网、
    气泡垫、无纺布、香梨纸、纸垫板、胶带、（备用）…），U「其他物料」＝总数－前面各列；表头在物料清单里找不到时标红；
    用户手打的「纸垫伴」改成「纸垫板」。客户名单只列有 领用/销售/退回 的单位。销售应收/已收/未收改成
    领用出库、销售出库、退回入库 三类按金额算（跟《03》口径一致）。冻结改到 C 列。
 2. 新表【发泡网销售汇总】：上块「库存与成本」（每种发泡网 期初/采购/销售/退回/结存、加权单价含期初、结存成本、
    本期销售成本），下块每客户一行（香梨网｜苹果网｜（备用）｜合计｜退回｜净数量｜销售金额｜已收｜未收）；带时段。
    「哪些物料算发泡网」就是下块表头的三个蓝格，改了全册（含物料成本接口）跟着变。
 3. 新表【其他包装物料汇总】：发泡网和筐子以外的包装物料，格局同 2；加权单价含期初。
 4. 新表【包装物料采购汇总】：每个供应商一行，物料列按表头求和＋（备用）＋其他，采购金额、已付、未付。
 5. 【筐子销售汇总】C～F 改成表头驱动（蓝色可改，默认 红/杂/新/一次性），红周转筐不再落到「其他」。
 6. 全部提取型汇总表（筐子库存/往来/采购/销售、自备筐、次果筐、押金、客户领用物料、物料库存结余）在第 4 行插「合计」，
    数据从第 5 行起；冻结、序号 ROW()-4、底部合计、条件格式、打印区，以及本册内部引用（实盘库存盘点、
    收入结算汇总、财务取数接口）全部跟着平移。顺手修：收入结算汇总 C21 指错列（M→O）、财务取数接口 G154 取了备注列
    （O→M）、物料库存结余「参考单价」把期初金额也算进加权（原来只看采购，期初物料库存金额都是 0）。
 7. 【筐子往来汇总】退回列 I～M 加上【次果筐+托盘明细】里业务类型＝「退回入库」的行（不含领用）；单位名单并入次果筐客户。
 8. 【周转筐出入库明细】在 N「装卸」后插 O「装筐费」（原 O～W → P～X，本册内所有引用、下拉、筛选、合并、隐藏列跟着移）：
    装筐费＝（入库数量＋退回数量）×装筐费单价；单价在【基础资料】U～V 列（统一单价＋按筐子类型分档，默认空＝不计费），
    期初、装卸方式空着或在「不计费的装卸方式」清单里（基础资料 U18～U23，默认 自卸、自装、自提、无）的不计。
    新表【装筐费接口】（K6）紧凑列出金额>0 的行，给《03》挂应付。
 9. 新表【物料成本接口】（K7）：B2 起始月（默认 2026-08-01），24 个月的 发泡网/其他包装物料 销售数量与销售成本，
    成本按「月末一次加权平均」（含期初）；隐藏计算网格在本表 K～Z 列。
10. 扩容：周转筐出入库明细 4～3003（3,000 行）、对账明细接口 4～3003（3,000 行，K2）；_自动清单 各池子跟着扩，
    逐行镜像全部改成按位置 INDEX 取（中间插行不再错位），补上周转筐/自备筐池子各少的 1 行，紧凑列表加「超过个数就不查」护栏。
11. 【主页】修两个指向不存在表名的链接（客户自备加工筐明细、次果筐+托盘明细），补上筐子采购汇总、对账明细接口，加新表入口；
    【使用说明】补本轮要点。
"""
import os
import re
import sys
import copy
import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from openpyxl.worksheet.cell_range import MultiCellRange
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.styles import Font, PatternFill, Alignment

from common0928 import (cell, title, note, section, header, widths, hide, period, dv_list,
                        F_TITLE, F_NOTE, F_HDR, F_SEC, F_LBL, F_IN, F_AUTO, F_AUTOB, F_TOT, F_TOTN, F_HELP,
                        FL_TITLE, FL_HDR, FL_SEC, FL_AUTO, FL_IN, FL_LBL, FL_TOT, FL_NOTE, FL_NONE,
                        AC, AL, AR, BOX, NOB, MONEY, QTY, INT, DATE, MONTH, fill)

# ════════════════════════════════════════════════════════════════════════════
#  一、公式引用改写引擎（插列、插行、区域加长、相对引用平移）
# ════════════════════════════════════════════════════════════════════════════
_REF = r'\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?'
_TOK = re.compile(
    r'(?P<str>"(?:[^"]|"")*")'
    r'|(?P<pre>(?:\'(?:[^\']|\'\')+\'|(?:\[\d+\])?[^\s!\'"()+\-*/&=<>,:;{}^%\[\]]+)!)'
    r'(?P<pref>' + _REF + r'|\$?\d+:\$?\d+|\$?[A-Z]{1,3}:\$?[A-Z]{1,3})'
    r'|(?P<bare>(?<![A-Za-z0-9_.$一-鿿\]\'!])' + _REF + r'(?![A-Za-z0-9_(一-鿿!]))')
_PART = re.compile(r'(\$?)([A-Z]{1,3})?(\$?)(\d+)?')


def _sheet_of(pre):
    """'xx'! / xx! → 表名；外部引用（[n]）返回 None"""
    s = pre[:-1]
    if s.startswith("'"):
        s = s[1:-1].replace("''", "'")
    if s.startswith('['):
        return None
    return s


def map_refs(f, host, fn):
    """对公式 f 里的每个单元格/区域引用调用 fn(表名, 引用文本) → 新引用文本。
       host＝公式所在表（不带表名前缀的引用算它的）；字符串常量、外部引用不动。"""
    if not isinstance(f, str):
        return f
    out, i = [], 0
    for m in _TOK.finditer(f):
        if m.group('str'):
            continue
        if m.group('pre'):
            sh = _sheet_of(m.group('pre'))
            if sh is None:
                continue
            old = m.group('pref')
            new = fn(sh, old)
            if new != old:
                out.append(f[i:m.start('pref')])
                out.append(new)
                i = m.end('pref')
        else:
            old = m.group('bare')
            new = fn(host, old)
            if new != old:
                out.append(f[i:m.start('bare')])
                out.append(new)
                i = m.end('bare')
    if not out:
        return f
    out.append(f[i:])
    return ''.join(out)


def _split(ref):
    """'$A$4:B9' → [(c$,col,r$,row), …]（列或行可能缺）"""
    res = []
    for p in ref.split(':'):
        m = _PART.fullmatch(p)
        res.append([m.group(1), m.group(2), m.group(3), m.group(4)])
    return res


def _join(parts):
    return ':'.join(f'{a}{b or ""}{c}{d or ""}' for a, b, c, d in parts)


def col_shifter(sheet, at, n=1):
    """在 sheet 的第 at 列前插入 n 列：引用到该表、列号>=at 的端点 +n"""
    def fn(sh, ref):
        if sh != sheet:
            return ref
        ps = _split(ref)
        for p in ps:
            if p[1] and CI(p[1]) >= at:
                p[1] = L(CI(p[1]) + n)
        return _join(ps)
    return fn


def row_shifter(sheet, at, n=1):
    """在 sheet 的第 at 行前插入 n 行：引用到该表、行号>=at 的端点 +n"""
    def fn(sh, ref):
        if sh != sheet:
            return ref
        ps = _split(ref)
        for p in ps:
            if p[3] and int(p[3]) >= at:
                p[3] = str(int(p[3]) + n)
        return _join(ps)
    return fn


def row_end_extender(sheet, start, old_end, new_end, absolute_only=True):
    """把引用到 sheet 的「start:old_end」区域加长到 new_end（只认绝对行号的结束端，滚动累计的 $H$4:$H9 这种不动）"""
    def fn(sh, ref):
        if sh != sheet or ':' not in ref:
            return ref
        ps = _split(ref)
        if len(ps) == 2 and ps[0][3] and ps[1][3] and int(ps[0][3]) == start and int(ps[1][3]) == old_end \
                and (ps[1][2] == '$' or not absolute_only):
            ps[1][3] = str(new_end)
        return _join(ps)
    return fn


def translator(dr, dc=0):
    """复制公式用：相对行/列号平移（绝对的不动），所有表的引用都算"""
    def fn(sh, ref):
        ps = _split(ref)
        for p in ps:
            if p[3] and p[2] != '$':
                p[3] = str(int(p[3]) + dr)
            if p[1] and dc and p[0] != '$':
                p[1] = L(CI(p[1]) + dc)
        return _join(ps)
    return fn


def tr(f, dr, dc=0):
    return map_refs(f, None, translator(dr, dc))


def shift_area(area, fn, sheet):
    """区域文字（可多个，用空格分隔，可带表名）整体改写"""
    out = []
    for a in str(area).split():
        if '!' in a:
            pre, ref = a.rsplit('!', 1)
            out.append(pre + '!' + fn(sheet, ref))
        else:
            out.append(fn(sheet, a))
    return ' '.join(out)


def map_workbook(wb, fn, only_if=None):
    """全工作簿的公式（单元格、数据验证、条件格式、定义名称）都过一遍 fn。
       only_if：公式里必须含这个子串（或是本表）才处理，提速用"""
    for ws in wb.worksheets:
        host = ws.title
        for c in ws._cells.values():
            v = c.value
            if isinstance(v, str) and v.startswith('=') and (only_if is None or only_if in v or host == only_if):
                nv = map_refs(v, host, fn)
                if nv != v:
                    c.value = nv
        for dv in ws.data_validations.dataValidation:
            for attr in ('formula1', 'formula2'):
                v = getattr(dv, attr)
                if v and (only_if is None or only_if in v or host == only_if):
                    setattr(dv, attr, map_refs(v, host, fn))
        for cf in ws.conditional_formatting:
            for rule in cf.rules:
                if rule.formula:
                    rule.formula = [map_refs(x, host, fn) for x in rule.formula]
        for name, dn in list(ws.defined_names.items()):
            if dn.attr_text:
                dn.attr_text = map_refs(dn.attr_text, host, fn)
    for name, dn in list(wb.defined_names.items()):
        if dn.attr_text:
            dn.attr_text = map_refs(dn.attr_text, None, fn)


def _geometry(ws, fn):
    """本表的几何属性（合并区、数据验证范围、条件格式范围、筛选、打印区）按 fn 改写"""
    t = ws.title
    merged = [str(m) for m in ws.merged_cells.ranges]
    for m in merged:
        ws.unmerge_cells(m)
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(shift_area(str(dv.sqref), fn, t))
    old = ws.conditional_formatting
    new = ConditionalFormattingList()
    for cf in old:
        rng = shift_area(str(cf.sqref), fn, t)
        for rule in cf.rules:
            new.add(rng, rule)
    ws.conditional_formatting = new
    if ws.auto_filter.ref:
        ws.auto_filter.ref = shift_area(ws.auto_filter.ref, fn, t)
    if ws.print_area:
        pa = [shift_area(a, fn, t).split('!')[-1] for a in str(ws.print_area).split(',')]
        ws.print_area = pa
    return merged


def insert_cols(wb, sheet, at, n=1, width=None):
    """整列插入：单元格（含样式、批注、超链接）、列宽/隐藏、合并、下拉、条件格式、筛选、打印区、全册公式都跟着走"""
    ws = wb[sheet]
    fn = col_shifter(sheet, at, n)
    merged = _geometry(ws, fn)
    ws.insert_cols(at, n)
    # 列宽：按 min/max 分组平移，跨插入点的组拆开
    groups = []
    for k, d in ws.column_dimensions.items():
        lo, hi = d.min or CI(k), d.max or CI(k)
        groups.append((lo, hi, d))
    ws.column_dimensions.clear()
    for lo, hi, d in groups:
        segs = []
        if hi < at:
            segs.append((lo, hi))
        elif lo >= at:
            segs.append((lo + n, min(hi + n, 16384)))
        else:
            segs += [(lo, at - 1), (at + n, min(hi + n, 16384))]
        for a, b in segs:
            nd = copy.copy(d)
            nd.index, nd.min, nd.max = L(a), a, b
            ws.column_dimensions[L(a)] = nd
    for k in range(at, at + n):
        d = ws.column_dimensions[L(k)]
        d.min = d.max = k
        d.width = width or 10
    if ws.freeze_panes:
        ws.freeze_panes = fn(sheet, ws.freeze_panes.replace('$', ''))
    map_workbook(wb, fn, only_if=sheet)
    for m in merged:
        ws.merge_cells(fn(sheet, m))


def insert_rows(wb, sheet, at, n=1):
    """整行插入（同上，行方向）"""
    ws = wb[sheet]
    fn = row_shifter(sheet, at, n)
    merged = _geometry(ws, fn)
    ws.insert_rows(at, n)
    rd = [(k, d) for k, d in ws.row_dimensions.items()]
    ws.row_dimensions.clear()
    for k, d in rd:
        nk = k + n if k >= at else k
        nd = copy.copy(d)
        nd.index = nk
        ws.row_dimensions[nk] = nd
    if ws.freeze_panes:
        ws.freeze_panes = fn(sheet, ws.freeze_panes.replace('$', ''))
    if ws.print_title_rows:
        pt = ws.print_title_rows
        ws.print_title_rows = fn(sheet, pt)
    map_workbook(wb, fn, only_if=sheet)
    for m in merged:
        ws.merge_cells(fn(sheet, m))
    # 本表被下移的公式里「ROW()-k」序号也跟着改
    for c in ws._cells.values():
        v = c.value
        if c.row >= at + n and isinstance(v, str) and v.startswith('=') and 'ROW()-' in v:
            c.value = re.sub(r'ROW\(\)-(\d+)', lambda m: f'ROW()-{int(m.group(1)) + n}', v)


# ════════════════════════════════════════════════════════════════════════════
#  二、常量
# ════════════════════════════════════════════════════════════════════════════
TZ = '周转筐出入库明细'
BZ = '包装物料出入库明细'
ZB = '客户自备加工筐明细'
CG = '次果筐+托盘明细'
AU = '_自动清单'
JK = '对账明细接口'
JC = '基础资料'
DY = '对接源_01水果筐数'
FPW = '发泡网销售汇总'
QTW = '其他包装物料汇总'
CGH = '包装物料采购汇总'
ZKF = '装筐费接口'
CBJ = '物料成本接口'

TZ_OLD1, TZ1 = 2004, 3003          # 周转筐出入库明细 数据区 4～3003（K4）
BZ1, ZB1, CG1 = 2003, 1204, 603     # 包装 / 自备 / 次果 数据区末行（不变）
JK_OLD1, JK1 = 1203, 3003          # 对账明细接口 4～3003（K2）
SUMMARY = [('物料库存结余', 54), ('筐子库存汇总', 64), ('筐子往来汇总', 254), ('筐子采购汇总', 204),
           ('筐子销售汇总', 204), ('自备筐汇总', 204), ('次果筐汇总', 204), ('押金汇总', 204),
           ('客户领用物料汇总', 204)]
FL_RED = PatternFill('solid', fgColor='FFC7CE')
F_RED = Font(color='9C0006', bold=True)
BAK = '（备用）'


def q(s):
    """公式里的表名：含 + 等符号时加单引号"""
    return f"'{s}'" if re.search(r'[^\w一-鿿]', s) else s


def R(sheet, col, r0, r1, absolute=True):
    d = '$' if absolute else ''
    return f'{q(sheet)}!{d}{col}{d}{r0}:{d}{col}{d}{r1}'


def bz(col):
    return R(BZ, col, 4, BZ1)


def tz(col):
    return R(TZ, col, 4, TZ1)


def cg(col):
    return R(CG, col, 4, CG1)


def au(col, r0=4, r1=BZ1):
    return f'{AU}!${col}${r0}:${col}${r1}'


def _copy_style(src, dst):
    if src.has_style:
        dst._style = copy.copy(src._style)


# ════════════════════════════════════════════════════════════════════════════
#  三、结构：插列、扩容、池子、对账明细接口、汇总表顶部合计
# ════════════════════════════════════════════════════════════════════════════
def tz_insert_fee_col(wb):
    """周转筐出入库明细：N（装卸）后插 O「装筐费」，原 O～W → P～X"""
    insert_cols(wb, TZ, 15, 1, width=10)
    ws = wb[TZ]
    ws['O3'] = '装筐费'
    _copy_style(ws['N3'], ws['O3'])


def tz_extend(wb):
    """周转筐出入库明细 数据区 4～2004 → 4～3003：公式、样式、行高、下拉、筛选、全册引用"""
    ws = wb[TZ]
    last = ws.max_column
    src_h = ws.row_dimensions[TZ_OLD1].height
    for col in range(1, last + 1):
        src = ws.cell(TZ_OLD1, col)
        f = src.value if isinstance(src.value, str) and src.value.startswith('=') else None
        for r in range(TZ_OLD1 + 1, TZ1 + 1):
            dst = ws.cell(r, col)
            if f:
                dst.value = tr(f, r - TZ_OLD1)
            _copy_style(src, dst)
    for r in range(TZ_OLD1 + 1, TZ1 + 1):
        if src_h:
            ws.row_dimensions[r].height = src_h
    for dv in ws.data_validations.dataValidation:
        dv.sqref = MultiCellRange(shift_area(str(dv.sqref), row_end_extender(TZ, 4, TZ_OLD1, TZ1, False), TZ))
    ws.auto_filter.ref = shift_area(ws.auto_filter.ref, row_end_extender(TZ, 3, TZ_OLD1, TZ1, False), TZ)
    map_workbook(wb, row_end_extender(TZ, 4, TZ_OLD1, TZ1), only_if=TZ)


def _pool3(ws, k, f, c, r0, r1, keyf, count=True):
    """三列去重池：k 候选值，f 首次出现标 1，c 累计编号；k1 放总个数"""
    for r in range(r0, r1 + 1):
        ws[f'{k}{r}'] = keyf(r)
        ws[f'{f}{r}'] = f'=IF({k}{r}="",0,IF(MATCH({k}{r},${k}${r0}:${k}${r1},0)=ROW()-{r0 - 1},1,0))'
        ws[f'{c}{r}'] = f'=N({c}{r - 1})+{f}{r}'
    if count:
        ws[f'{k}1'] = f'=MAX(${c}${r0}:${c}${r1})'


def _compact(ws, col, r1, k, c, kr1, cnt_cell, src=None):
    """紧凑列表（第 4 行起），超过个数就不查"""
    src = src or k
    for r in range(4, r1 + 1):
        ws[f'{col}{r}'] = (f'=IF(ROW()-3>N({cnt_cell}),"",IFERROR(INDEX(${src}$4:${src}${kr1},'
                           f'MATCH(ROW()-3,${c}$4:${c}${kr1},0)),""))')


def _clear_col(ws, col, r0, r1):
    for r in range(r0, r1 + 1):
        if (r, CI(col)) in ws._cells:
            ws.cell(r, CI(col)).value = None


def _ix(sheet, col, r0, r1, off):
    """按位置取：INDEX(表!$X$r0:$X$r1,ROW()-off)"""
    return f'INDEX({R(sheet, col, r0, r1)},ROW()-{off})'


def pools(wb):
    """_自动清单：逐行镜像改成按位置 INDEX；周转筐池子扩到 3,000 行、自备池子补第 1204 行；紧凑列表加护栏"""
    ws = wb[AU]
    # 包装物料：物料种类 A/D/E（规格、单位），物料对方 K
    for r in range(4, BZ1 + 1):
        e = _ix(BZ, 'E', 4, BZ1, 3)
        ws[f'A{r}'] = f'=IF({e}="","",{e})'
        ws[f'D{r}'] = f'=IF({e}="","",{_ix(BZ, "F", 4, BZ1, 3)})'
        ws[f'E{r}'] = f'=IF({e}="","",{_ix(BZ, "G", 4, BZ1, 3)})'
        d = _ix(BZ, 'D', 4, BZ1, 3)
        ws[f'K{r}'] = f'=IF({d}="","",{d})'
    _compact(ws, 'G', 63, 'A', 'C', BZ1, '$A$1')
    _compact(ws, 'H', 63, 'A', 'C', BZ1, '$A$1', src='D')
    _compact(ws, 'I', 63, 'A', 'C', BZ1, '$A$1', src='E')
    _compact(ws, 'O', 203, 'K', 'M', BZ1, '$K$1')
    # 周转筐：类型 Q/R/S、对方 W/X/Y、采购供应商 CS/CT/CU（4～3003）
    for col in ('Q', 'R', 'S', 'W', 'X', 'Y', 'CS', 'CT', 'CU'):
        _clear_col(ws, col, 4, TZ1)
    dz, cz, fz = (_ix(TZ, x, 4, TZ1, 3) for x in 'DCF')
    _pool3(ws, 'Q', 'R', 'S', 4, TZ1, lambda r: f'=IF({dz}="","",{dz})')
    ws['Q1'] = f'=MAX($S$4:$S${TZ1})'
    _pool3(ws, 'W', 'X', 'Y', 4, TZ1, lambda r: f'=IF({cz}="","",{cz})')
    ws['W1'] = f'=MAX($Y$4:$Y${TZ1})'
    _pool3(ws, 'CS', 'CT', 'CU', 4, TZ1, lambda r: f'=IF(OR({cz}="",{fz}<>"采购入库"),"",{cz})')
    ws['CS1'] = f'=MAX($CU$4:$CU${TZ1})'
    _compact(ws, 'U', 43, 'Q', 'S', TZ1, '$Q$1')
    _compact(ws, 'AA', 203, 'W', 'Y', TZ1, '$W$1')
    _compact(ws, 'CW', 203, 'CS', 'CU', TZ1, '$CS$1')
    # 客户自备加工筐：类型 AC（4～1204），客户 AI（4～1204 自备明细 ＋ 1205～1504 对接源_01 自备段）
    for col in ('AC', 'AD', 'AE', 'AI', 'AJ', 'AK'):
        _clear_col(ws, col, 4, 1504)
    zc, ze = _ix(ZB, 'C', 4, ZB1, 3), _ix(ZB, 'E', 4, ZB1, 3)
    _pool3(ws, 'AC', 'AD', 'AE', 4, ZB1, lambda r: f'=IF({zc}="","",{zc})')
    ws['AC1'] = f'=MAX($AE$4:$AE${ZB1})'

    def ai(r):
        if r <= ZB1:
            return f'=IF({ze}="","",{ze})'
        b, c = _ix(DY, 'B', 5, 304, ZB1), _ix(DY, 'C', 5, 304, ZB1)
        return f'=IF(OR({b}="",LEFT({c},2)<>"自备"),"",{b})'
    _pool3(ws, 'AI', 'AJ', 'AK', 4, ZB1 + 300, ai)
    ws['AI1'] = f'=MAX($AK$4:$AK${ZB1 + 300})'
    _compact(ws, 'AG', 43, 'AC', 'AE', ZB1, '$AC$1')
    _compact(ws, 'AM', 203, 'AI', 'AK', ZB1 + 300, '$AI$1')
    # 次果筐：类型 AO、客户 AU（4～603）
    for r in range(4, CG1 + 1):
        d, c = _ix(CG, 'D', 4, CG1, 3), _ix(CG, 'C', 4, CG1, 3)
        ws[f'AO{r}'] = f'=IF({d}="","",{d})'
        ws[f'AU{r}'] = f'=IF({c}="","",{c})'
    _compact(ws, 'AS', 43, 'AO', 'AQ', CG1, '$AO$1')
    _compact(ws, 'AY', 203, 'AU', 'AW', CG1, '$AU$1')
    # 对接源_01水果筐数（5～304 → 池 4～303）
    for r in range(4, 304):
        b, c = _ix(DY, 'B', 5, 304, 3), _ix(DY, 'C', 5, 304, 3)
        ws[f'BA{r}'] = f'=IF({b}="","",{b})'
        ws[f'BG{r}'] = f'=IF({c}="","",{c})'
    _compact(ws, 'BE', 303, 'BA', 'BC', 303, '$BA$1')
    _compact(ws, 'BK', 63, 'BG', 'BI', 303, '$BG$1')
    _compact(ws, 'BQ', 63, 'BM', 'BO', 183, '$BM$1')
    _compact(ws, 'BW', 253, 'BS', 'BU', 1103, '$BS$1')
    # 筐子往来客户 BY：周转筐对方 AA（4～203）＋01 册筐客户 BE（204～503）＋次果筐客户 AY（504～703）
    for col in ('BY', 'BZ', 'CA'):
        _clear_col(ws, col, 4, 703)

    def by(r):
        if r <= 203:
            return f'=IF($AA{r}="","",$AA{r})'
        if r <= 503:
            return f'=IF($BE{r - 200}="","",$BE{r - 200})'
        return f'=IF($AY{r - 500}="","",$AY{r - 500})'
    _pool3(ws, 'BY', 'BZ', 'CA', 4, 703, by)
    ws['BY1'] = '=MAX($CA$4:$CA$703)'
    ws['BY3'] = '筐子往来客户（含次果筐客户）'
    _compact(ws, 'CC', 253, 'BY', 'CA', 703, '$BY$1')
    # 基础资料镜像：单位 CE、客户/供应商 CI
    for r in range(4, 54):
        p = _ix(JC, 'P', 4, 53, 3)
        ws[f'CE{r}'] = f'=IF({p}="","",{p})'
    for r in range(4, 204):          # 4～103＝基础资料 D「客户/领取人」，104～203＝E「供应商」
        d = _ix(JC, 'D', 4, 103, 3) if r <= 103 else _ix(JC, 'E', 4, 103, 103)
        ws[f'CI{r}'] = f'=IF({d}="","",{d})'
    # 对账命中 CP/CQ：包装 4～2003｜周转筐 2004～5003｜自备 5004～6204｜次果 6205～6804
    _clear_col(ws, 'CP', 4, 6804)
    _clear_col(ws, 'CQ', 4, 6804)
    segs = [(4, BZ1, BZ, 'D', BZ1), (BZ1 + 1, BZ1 + 3000, TZ, 'C', TZ1),
            (BZ1 + 3001, BZ1 + 3001 + 1200, ZB, 'E', ZB1), (BZ1 + 3001 + 1201, BZ1 + 3001 + 1201 + 599, CG, 'C', CG1)]
    for a, b, sh, col, end in segs:
        for r in range(a, b + 1):
            ws[f'CP{r}'] = f'=IF(INDEX({R(sh, col, 4, end)},ROW()-{a - 1})="",0,1)'
            ws[f'CQ{r}'] = f'=N(CQ{r - 1})+$CP{r}'
    ws['CP1'] = '=MAX($CQ$4:$CQ$6804)'


_SEG_LE = {'2000': '2000', '4000': '5000', '5200': '6201', '5800': '6801'}
_SEG_MINUS = {'0': '0', '2000': '2000', '4000': '5000', '5200': '6201'}


def _reseg(f):
    f = re.sub(r'(\$Z\d+)<=(\d+)', lambda m: m.group(1) + '<=' + _SEG_LE[m.group(2)], f)
    f = re.sub(r'(\$Z\d+)-(\d+)(?!\d)', lambda m: m.group(1) + '-' + _SEG_MINUS[m.group(2)], f)
    f = re.sub(r'(\$Z\d+)>(\d+)', lambda m: m.group(1) + '>' + {'4000': '5000'}[m.group(2)], f)
    return f


def jk_extend(wb):
    """对账明细接口 4～1203 → 4～3003（K2），分段阈值改成 2000/5000/6201/6801，Z 列加护栏"""
    ws = wb[JK]
    tpl = {}
    for col in list('ABCDEFGHIJKLMNOPQRS'):
        v = ws[f'{col}4'].value
        tpl[col] = _reseg(v) if isinstance(v, str) and v.startswith('=') else None
    tpl['Z'] = (f'=IF(ROW()-3>N({AU}!$CP$1),"",IFERROR(MATCH(ROW()-3,{AU}!$CQ$4:$CQ$6804,0),""))')
    src_h = ws.row_dimensions[JK_OLD1].height
    for col, f in tpl.items():
        st = ws[f'{col}{JK_OLD1}']
        for r in range(4, JK1 + 1):
            c = ws[f'{col}{r}']
            c.value = tr(f, r - 4) if f else None
            if r > JK_OLD1:
                _copy_style(st, c)
    if src_h:
        for r in range(JK_OLD1 + 1, JK1 + 1):
            ws.row_dimensions[r].height = src_h
    n = f'N({AU}!$CP$1)'
    ws['A2'] = (f'=IF({n}>3000,"★★ 注意：业务笔数已有 "&{n}&" 个，超过本表 3000 个的显示上限，请联系维护人员扩容（《04》也要一起扩）",'
                f'"★《04》按固定位置引用本表 A3:S3003，请不要插入/删除行列或改表名。本表把包装物料、周转筐、客户自备筐、'
                f'次果筐托盘四张明细的每一笔业务规范化成同一种格式，并按业务类别成组排列　【业务笔数 "&{n}&"/3000】　'
                f'★「发出筐数」只取《周转筐出入库明细》的「出库数量」，不含「销售数量」——卖断的筐子不是往来筐，'
                f'它走金额线，在《04 物料与筐子销售对账单》里对账。")')
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if rule.formula:
                rule.formula = [x.replace('>1200)', '>3000)') for x in rule.formula]


def top_totals(wb):
    """各汇总表第 4 行插「合计」，数据从第 5 行起（底部合计保留）"""
    for name, tot in SUMMARY:
        insert_rows(wb, name, 4, 1)
        ws = wb[name]
        bot = tot + 1
        for col in range(1, ws.max_column + 1):
            s = ws.cell(bot, col)
            d = ws.cell(4, col)
            _copy_style(s, d)
            if isinstance(s.value, str) and s.value.startswith('='):
                d.value = s.value
        ws['B4'] = '合计'
        h = ws.row_dimensions[bot].height
        ws.row_dimensions[4].height = h or 20


def structure(wb):
    tz_insert_fee_col(wb)
    tz_extend(wb)
    pools(wb)
    jk_extend(wb)
    top_totals(wb)


# ════════════════════════════════════════════════════════════════════════════
#  四、装筐费（周转筐 O 列、基础资料参数、装筐费接口）
# ════════════════════════════════════════════════════════════════════════════
FEE_FREE = ['自卸', '自装', '自提', '无', None, None]     # 默认不计装筐费的装卸方式（基础资料 U18～U23）
FEE_TYPES = ['公司一次性筐', '公司杂周转筐', '公司新周转筐', '公司白周转筐', '公司红周转筐',
             '自备一次性筐', '自备杂周转筐', '大托盘', None, None]


def fee_param(wb):
    """基础资料 U～W：装筐费单价（统一单价＋按筐子类型分档），默认空＝不计费"""
    ws = wb[JC]
    hs = ws['R3']
    for ref, v in (('U3', '装筐费单价'), ('V3', '元/个'), ('W3', '说明')):
        ws[ref] = v
        _copy_style(hs, ws[ref])
    cell(ws, 'U4', '统一单价', F_LBL, FL_LBL, AC)
    cell(ws, 'V4', None, F_IN, FL_IN, AC, fmt='0.00##')
    cell(ws, 'U5', '按筐子类型分档 ↓', F_LBL, FL_LBL, AC)
    cell(ws, 'V5', '单价（空＝用统一价）', F_LBL, FL_LBL, AC)
    for i, t in enumerate(FEE_TYPES):
        cell(ws, f'U{6 + i}', t, F_IN, FL_IN, AC)
        cell(ws, f'V{6 + i}', None, F_IN, FL_IN, AC, fmt='0.00##')
    cell(ws, 'U17', '不计费的装卸方式 ↓', F_LBL, FL_LBL, AC)
    cell(ws, 'V17', '（可增删）', F_LBL, FL_LBL, AC)
    for i, t in enumerate(FEE_FREE):
        cell(ws, f'U{18 + i}', t, F_IN, FL_IN, AC)
        cell(ws, f'V{18 + i}', None, F_NOTE, FL_NONE, AC)
    ws.merge_cells('W4:W15')
    cell(ws, 'W4', '★ 装筐费＝（入库数量＋退回数量）× 单价，自动算在【周转筐出入库明细】O 列，付给「装卸」列的装卸方'
                   '（如 依力），《03》挂应付账款。\n· 两个单价都空着＝不计费（默认）。\n· 某种筐子单独填了分档单价就用分档价，'
                   '否则用统一单价。\n· 期初库存行、装卸方式空着或写在 U18～U23「不计费的装卸方式」里的行（默认 自卸、自装、自提、无）不计；'
                   '「筐厂」卸货要不要计，定了以后把它加进去或删掉就行。\n· 类型名要和 B 列筐子类型一字不差。',
         F_NOTE, FL_NOTE, AL)
    for r in range(5, 16):
        ws[f'W{r}'].border = BOX
    widths(ws, {'U': 18, 'V': 14, 'W': 46})


def fee_price(r):
    s = f'SUMIFS({JC}!$V$6:$V$15,{JC}!$U$6:$U$15,$D{r})'
    return f'IF({s}>0,{s},N({JC}!$V$4))'


def tz_fee(wb):
    ws = wb[TZ]
    st = ws['Q4']
    for r in range(4, TZ1 + 1):
        p = fee_price(r)
        ws[f'O{r}'] = (f'=IF(OR($D{r}="",$F{r}="期初库存",$N{r}="",COUNTIF({JC}!$U$18:$U$23,$N{r})>0,'
                       f'N($H{r})+ABS(N($I{r}))=0),"",'
                       f'IF({p}=0,"",ROUND((N($H{r})+ABS(N($I{r})))*{p},2)))')
        _copy_style(ws[f'Q{r}'] if r > 4 else st, ws[f'O{r}'])
        ws[f'O{r}'].number_format = MONEY
    a2 = ws['A2'].value or ''
    if '装筐费' not in a2:
        ws['A2'] = (a2 + '　★ O 列「装筐费」自动算：（入库数量＋退回数量）×装筐费单价，单价在【基础资料】U～V 列填'
                    '（留空＝不计费）；期初、装卸方式空着或属于【基础资料】U18～U23「不计费的装卸方式」（默认 自卸、自装、自提、无）的不计。'
                    '本表可录到第 3003 行，请在最后一行下面接着录，不要在中间插行。')


def build_fee_iface(wb, idx):
    """新表【装筐费接口】（K6）：行 4～503，紧凑列出装筐费>0 的行"""
    ws = wb.create_sheet(ZKF, idx)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = 'C55A11'
    title(ws, '装 筐 费 接 口（全自动 · 供《03 财务账套》挂应付，请勿改动结构）', 'A', 'I')
    cell(ws, 'K1', '笔数', F_LBL, FL_LBL, AC)
    cell(ws, 'L1', f'=N({AU}!$ED$1)', F_AUTOB, FL_AUTO, AC, fmt=INT)
    note(ws, f'★ 按【周转筐出入库明细】O 列「装筐费」逐笔列出金额>0 的行（第 4 行起连续排列，最多 500 笔）。'
             f'付款对象＝该行「装卸」列。装筐费单价在【基础资料】U～V 列设，默认空＝不计费，所以本表平时是空的。', 'A', 'I', 2, 36)
    header(ws, 3, 'A', ['序号', '日期', '付款对象\n（装卸方）', '筐子类型', '计费数量', '装筐费单价', '装筐费金额', '原表行号', '摘要'])
    for r in range(4, 504):
        h = f'$H{r}'
        ix = lambda col: f'INDEX({tz(col)},{h}-3)'
        vals = {
            'A': f'=IF({h}="","",ROW()-3)',
            'B': f'=IF({h}="","",IF({ix("B")}="","",{ix("B")}))',
            'C': f'=IF({h}="","",{ix("N")}&"")',
            'D': f'=IF({h}="","",{ix("D")}&"")',
            'E': f'=IF({h}="","",N({ix("H")})+ABS(N({ix("I")})))',
            'F': f'=IF(OR({h}="",N($E{r})=0),"",ROUND(N($G{r})/N($E{r}),4))',
            'G': f'=IF({h}="","",N({ix("O")}))',
            'H': f'=IF(ROW()-3>N({AU}!$ED$1),"",IFERROR(MATCH(ROW()-3,{AU}!$EE$4:$EE${TZ1},0)+3,""))',
            'I': (f'=IF({h}="","","装筐费 "&{ix("F")}&" "&{ix("C")}&" "&$D{r}&" "&$E{r}&"个"'
                  f'&IF({ix("E")}="",""," "&{ix("E")}))'),
        }
        fmts = {'B': DATE, 'E': INT, 'F': '0.00##', 'G': MONEY, 'H': '0'}
        for col, f in vals.items():
            cell(ws, f'{col}{r}', f, F_AUTO, FL_AUTO, AL if col in 'CI' else AC, fmt=fmts.get(col))
    widths(ws, {'A': 6, 'B': 11, 'C': 12, 'D': 14, 'E': 10, 'F': 11, 'G': 12, 'H': 9, 'I': 46, 'K': 7, 'L': 8})
    ws.freeze_panes = 'C4'


# ════════════════════════════════════════════════════════════════════════════
#  五、新池子（_自动清单 CY 列以后）
# ════════════════════════════════════════════════════════════════════════════
FPW_HDR = f'{FPW}!$C$12:$E$12'      # 「哪些物料算发泡网」＝发泡网销售汇总下块表头三个蓝格


def _cat(x):
    """物料 x 的类别：发泡网 / 筐 / 其他"""
    return (f'IF({x}="","",IF(COUNTIF({FPW_HDR},{x})>0,"发泡网",IF(OR(COUNTIF({JC}!$B$4:$B$103,{x})>0,'
            f'COUNTIF({JC}!$C$4:$C$103,{x})>0,RIGHT({x},1)="筐"),"筐","其他")))')


def new_pools(wb):
    ws = wb[AU]
    hdrs = {'CY': '销售类(领用/销售/退回)', 'CZ': '物料类别', 'DA': '领用+销售出库数', 'DB': '退回入库数',
            'DC': '销售金额(三类)', 'DD': '其中已收讫',
            'DF': '领用客户', 'DG': '首次', 'DH': '累计', 'DJ': '领用客户(紧凑)',
            'DL': '发泡网客户', 'DM': '首次', 'DN': '累计', 'DP': '发泡网客户(紧凑)',
            'DR': '其他物料客户', 'DS': '首次', 'DT': '累计', 'DV': '其他物料客户(紧凑)',
            'DX': '包装采购供应商', 'DY': '首次', 'DZ': '累计', 'EB': '包装采购供应商(紧凑)',
            'ED': '装筐费>0', 'EE': '累计', 'EG': '其他物料', 'EH': '首次', 'EI': '累计', 'EK': '其他物料(紧凑)'}
    for k, v in hdrs.items():
        ws[f'{k}3'] = v
    for r in range(4, BZ1 + 1):
        h, e, d = (_ix(BZ, x, 4, BZ1, 3) for x in 'HED')
        ws[f'CY{r}'] = f'=IF(OR({h}="领用出库",{h}="销售出库",{h}="退回入库"),1,0)'
        ws[f'CZ{r}'] = '=' + _cat(e)
        ws[f'DA{r}'] = f'=IF(OR({h}="领用出库",{h}="销售出库"),N({_ix(BZ, "J", 4, BZ1, 3)}),0)'
        ws[f'DB{r}'] = f'=IF({h}="退回入库",N({_ix(BZ, "K", 4, BZ1, 3)}),0)'
        ws[f'DC{r}'] = f'=IF($CY{r}=1,N({_ix(BZ, "N", 4, BZ1, 3)}),0)'
        ws[f'DD{r}'] = f'=IF(AND($CY{r}=1,{_ix(BZ, "P", 4, BZ1, 3)}="已收讫"),N({_ix(BZ, "N", 4, BZ1, 3)}),0)'
    d = _ix(BZ, 'D', 4, BZ1, 3)
    h = _ix(BZ, 'H', 4, BZ1, 3)
    _pool3(ws, 'DF', 'DG', 'DH', 4, BZ1, lambda r: f'=IF(AND($CY{r}=1,{d}<>""),{d},"")')
    _pool3(ws, 'DL', 'DM', 'DN', 4, BZ1, lambda r: f'=IF(AND($CY{r}=1,$CZ{r}="发泡网",{d}<>""),{d},"")')
    _pool3(ws, 'DR', 'DS', 'DT', 4, BZ1, lambda r: f'=IF(AND($CY{r}=1,$CZ{r}="其他",{d}<>""),{d},"")')
    _pool3(ws, 'DX', 'DY', 'DZ', 4, BZ1, lambda r: f'=IF(AND({h}="采购入库",{d}<>""),{d},"")')
    _compact(ws, 'DJ', 203, 'DF', 'DH', BZ1, '$DF$1')
    _compact(ws, 'DP', 203, 'DL', 'DN', BZ1, '$DL$1')
    _compact(ws, 'DV', 203, 'DR', 'DT', BZ1, '$DR$1')
    _compact(ws, 'EB', 203, 'DX', 'DZ', BZ1, '$DX$1')
    # 装筐费>0 的周转筐行
    for r in range(4, TZ1 + 1):
        ws[f'ED{r}'] = f'=IF(N({_ix(TZ, "O", 4, TZ1, 3)})>0,1,0)'
        ws[f'EE{r}'] = f'=N(EE{r - 1})+ED{r}'
    ws['ED1'] = f'=MAX($EE$4:$EE${TZ1})'
    # 其他包装物料（物料种类里不是发泡网、不是筐的）
    _pool3(ws, 'EG', 'EH', 'EI', 4, 63, lambda r: f'=IF({_cat(f"$G{r}")}="其他",$G{r},"")')
    _compact(ws, 'EK', 63, 'EG', 'EI', 63, '$EG$1')
    for k in hdrs:
        for r in (1, 3):
            ws[f'{k}{r}'].font = F_HELP


# ════════════════════════════════════════════════════════════════════════════
#  六、改老表
# ════════════════════════════════════════════════════════════════════════════
def _win(s, e):
    return f'{bz("B")},">="&{s},{bz("B")},"<="&{e}'


def _red_header_cf(ws, rng, list_rng):
    tl = rng.split(':')[0]
    ws.conditional_formatting.add(rng, FormulaRule(
        formula=[f'AND({tl}<>"",LEFT({tl},1)<>"（",LEFT({tl},1)<>"(",COUNTIF({list_rng},{tl})=0)'],
        fill=FL_RED, font=F_RED))


def _hdr_blank(c):
    return f'OR({c}="",LEFT({c},1)="（",LEFT({c},1)="(")'


def fixes(wb):
    wb['收入结算汇总']['C21'] = '=自备筐汇总!$O$205'          # 原来取 M（自备一次性筐剩余），应为 O 剩余合计
    wb['财务取数接口']['G154'] = '=ROUND(物料库存结余!$M$55,2)'  # 原来取 O（备注列）
    # 物料库存结余：参考单价改成「(期初金额＋采购金额)÷(期初数量＋采购数量)」
    ws = wb['物料库存结余']
    w = f'{bz("B")},">="&$S$1,{bz("B")},"<="&$S$2'
    for r in range(5, 55):
        a = lambda col, t: f'SUMIFS({bz(col)},{bz("E")},$B{r},{bz("H")},"{t}",{w})'
        ws[f'E{r}'] = (f'=IF($B{r}="","",IFERROR(ROUND(({a("N", "期初库存")}+{a("N", "采购入库")})/'
                       f'({a("I", "期初库存")}+{a("I", "采购入库")}),4),0))')
    ws['E3'] = '加权单价\n(含期初)'


def collect_material(wb):
    """客户领用物料汇总：表头驱动、应收三类、只列有领用/销售/退回的单位、冻结 C 列"""
    ws = wb['客户领用物料汇总']
    names = ['香梨网', '苹果网', '气泡垫', '无纺布', '香梨纸', '纸垫板', '胶带'] + [BAK] * 11
    cols = [L(i) for i in range(3, 21)]          # C～T
    for col, t in zip(cols, names):
        cell(ws, f'{col}3', t, F_IN, FL_IN, AC)
    _red_header_cf(ws, 'C3:T3', f'{JC}!$A$4:$A$103')
    dv_list(ws, 'C3:T3', '物料种类表')
    win = f'{bz("B")},">="&$AF$1,{bz("B")},"<="&$AF$2'
    for r in range(5, 205):
        ws[f'B{r}'] = f'=IFERROR(INDEX({AU}!$DJ$4:$DJ$203,ROW()-4),"")'
        for col in cols:
            ws[f'{col}{r}'] = (f'=IF(OR($B{r}="",{_hdr_blank(col + "$3")}),"",SUMIFS({bz("J")},{bz("D")},$B{r},'
                               f'{bz("E")},{col}$3,{win}))')
        ws[f'Y{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({au("DC")},{bz("D")},$B{r},{win}),2))'
        ws[f'Z{r}'] = f'=IF($B{r}="","",ROUND(SUMIFS({au("DD")},{bz("D")},$B{r},{win}),2))'
    ws['Y3'] = '应收金额\n(领用/销售/退回)'
    ws['A2'] = ('★ 全自动。C～T 列按第 3 行表头（蓝格，可改成别的物料名）统计各客户的出库数量，表头在物料清单里找不到会标红；'
                'U「其他物料」＝没列出来的物料。应收金额＝领用出库＋销售出库＋退回入库（退回冲减）的合计金额，跟《03》口径一致；'
                '已收＝其中收付款状态为「已收讫」的部分。填右边【开始日期／结束日期】按区间统计（留空＝不限）。')
    ws.freeze_panes = 'C5'
    ws.sheet_view.selection[0].activeCell = 'C5'
    ws.sheet_view.selection[0].sqref = 'C5'
    new = ConditionalFormattingList()
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            if str(cf.sqref) == 'A2':
                rule.formula = [f'(N({AU}!$DF$1)>200)']
            new.add(str(cf.sqref), rule)
    ws.conditional_formatting = new


def basket_sales(wb):
    """筐子销售汇总：C～F 改成表头驱动（蓝格可改）"""
    ws = wb['筐子销售汇总']
    for col, t in zip('CDEF', ['公司红周转筐', '公司杂周转筐', '公司新周转筐', '公司一次性筐']):
        cell(ws, f'{col}3', t, F_IN, FL_IN, AC)
    _red_header_cf(ws, 'C3:F3', f'{JC}!$B$4:$B$103')
    dv_list(ws, 'C3:F3', '筐子类型表')
    for r in range(5, 205):
        for col in 'CDEF':
            ws[f'{col}{r}'] = (f'=IF(OR($B{r}="",{_hdr_blank(col + "$3")}),"",SUMIFS({tz("K")},{tz("C")},$B{r},'
                               f'{tz("D")},{col}$3,{tz("B")},">="&$P$1,{tz("B")},"<="&$P$2))')
    a2 = ws['A2'].value or ''
    if '表头' not in a2:
        ws['A2'] = a2 + ' C～F 列按第 3 行表头（蓝格，可改成别的筐子类型）统计，没列出来的类型进「其他」。'


def basket_flow(wb):
    """筐子往来汇总：退回列 I～M 加上次果筐+托盘明细里「退回入库」的行"""
    ws = wb['筐子往来汇总']
    for r in range(5, 255):
        for i, col in enumerate('IJKLM'):
            t = f'{JC}!$B${4 + i}'
            ws[f'{col}{r}'] = (f'=IF(OR($B{r}="",{t}=""),"",-SUMIFS({tz("I")},{tz("C")},$B{r},{tz("D")},{t},'
                               f'{tz("B")},">="&$BC$1,{tz("B")},"<="&$BC$2)'
                               f'-SUMIFS({cg("G")},{cg("C")},$B{r},{cg("D")},{t},{cg("F")},"退回入库",'
                               f'{cg("B")},">="&$BC$1,{cg("B")},"<="&$BC$2))')
    ws['N3'] = '退回合计\n(含次果筐退回)'
    a2 = ws['A2'].value or ''
    if '次果筐' not in a2:
        ws['A2'] = a2 + ' 退回列 I～M 含【次果筐+托盘明细】里「退回入库」的筐（次果筐领用不算发出）。'


# ════════════════════════════════════════════════════════════════════════════
#  七、新表
# ════════════════════════════════════════════════════════════════════════════
def _new_sheet(wb, name, idx, tab='548235'):
    ws = wb.create_sheet(name, idx)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    return ws


def _neg_fmt(wb):
    return wb['客户领用物料汇总']['W5'].number_format or '[Red]\\-#,##0;[Red]#,##0;\\-'


def _stock_block(ws, r0, r1, mat_formula, p):
    """库存与成本块：r0..r1 每行一种物料（A 列＝物料名公式），r1+1 合计"""
    S, E = p['start'], p['end']
    win = ',' + _win(S, E)
    pre = ',' + bz('B') + ',"<"&' + S
    H = bz('H')
    c_qc, c_nqc, c_cg = ',%s,"期初库存"' % H, ',%s,"<>期初库存"' % H, ',%s,"采购入库"' % H
    header(ws, r0 - 1, 'A', ['物料', '期初数量', '期初金额', '采购数量', '采购金额', '销售(领用)\n数量', '客户退回\n数量',
                             '其他减少\n(自用/报损等)', '结存数量', '加权单价\n(含期初)', '结存成本', '本期销售成本'])
    for r in range(r0, r1 + 1):
        m = f'$A{r}'

        def s(col, *extra):
            return 'SUMIFS(%s,%s,%s%s)' % (bz(col), bz('E'), m, ''.join(extra))
        f = {
            'A': mat_formula(r),
            'B': f'=IF({m}="","",{s("I", c_qc)}+{s("I", c_nqc, pre)}-{s("J", pre)}+{s("K", pre)})',
            'C': (f'=IF({m}="","",IF(N($B{r})=0,0,ROUND(N($B{r})*IFERROR(({s("N", c_qc)}+{s("N", c_cg, pre)})/'
                  f'({s("I", c_qc)}+{s("I", c_cg, pre)}),0),2)))'),
            'D': f'=IF({m}="","",{s("I", c_cg, win)})',
            'E': f'=IF({m}="","",ROUND({s("N", c_cg, win)},2))',
            'F': f'=IF({m}="","",SUMIFS({au("DA")},{bz("E")},{m}{win}))',
            'G': f'=IF({m}="","",SUMIFS({au("DB")},{bz("E")},{m}{win}))',
            'H': f'=IF({m}="","",{s("J", win)}-N($F{r}))',
            'I': f'=IF({m}="","",N($B{r})+N($D{r})-N($F{r})+N($G{r})-N($H{r}))',
            'J': f'=IF({m}="","",IFERROR(ROUND((N($C{r})+N($E{r}))/(N($B{r})+N($D{r})),4),0))',
            'K': f'=IF({m}="","",ROUND(N($I{r})*N($J{r}),2))',
            'L': f'=IF({m}="","",ROUND((N($F{r})-N($G{r}))*N($J{r}),2))',
        }
        fm = {'C': MONEY, 'E': MONEY, 'J': '#,##0.0000;[Red]\\-#,##0.0000;\\-', 'K': MONEY, 'L': MONEY}
        for col, v in f.items():
            cell(ws, f'{col}{r}', v, F_AUTO, FL_AUTO, AC, fmt=fm.get(col, QTY))
    t = r1 + 1
    cell(ws, f'A{t}', '合计', F_TOT, FL_TOT, AC)
    for col in 'BCDEFGHIKL':
        cell(ws, f'{col}{t}', f'=ROUND(SUM({col}{r0}:{col}{r1}),2)', F_TOTN, FL_TOT, AC,
             fmt=MONEY if col in 'CEKL' else QTY)
    cell(ws, f'J{t}', None, F_TOTN, FL_TOT, AC)


def build_fpw(wb, idx):
    ws = _new_sheet(wb, FPW, idx)
    title(ws, '发 泡 网 销 售 汇 总（香梨网＋苹果网 · 库存成本＋按客户 · 全自动）', 'A', 'L')
    p = period(ws, row=2, col='A', open_ended=True, year=None)
    note(ws, '★ 全自动。「哪些物料算发泡网」就是第 12 行 C～E 的蓝格（默认 香梨网、苹果网，第三格备用，可改），'
             '《物料成本接口》也按这几格分类。上块：期初＝起始日以前的结存（含建账期初），加权单价＝(期初金额＋采购金额)÷(期初数量＋采购数量)，'
             '结存成本＝结存数量×加权单价，本期销售成本＝(销售−退回)×加权单价。下块：每个客户的领用/销售数量、客户退回、销售金额'
             '（领用＋销售−退回，跟《03》口径一致）和已收（收付款状态＝已收讫）。第 2 行填年份或起止日按区间统计，都空＝不限。',
         'A', 'L', 3, 48)
    section(ws, '一、库存与成本（按物料）', 'A', 'L', 4)
    mats = ['C', 'D', 'E']
    _stock_block(ws, 6, 8, lambda r: '=IF(%s,"",%s$12)' % (_hdr_blank(mats[r - 6] + '$12'), mats[r - 6]), p)
    section(ws, '二、按客户（领用出库＋销售出库，客户退回冲减）', 'A', 'L', 11)
    header(ws, 12, 'A', ['序号', '客户', '香梨网', '苹果网', BAK, '合计数量', '客户退回\n(显示为负)', '净数量',
                          '销售金额', '已收金额', '未收金额', '备注'])
    for col in mats:
        cell(ws, f'{col}12', ws[f'{col}12'].value, F_IN, FL_IN, AC)
    _red_header_cf(ws, 'C12:E12', f'{JC}!$A$4:$A$103')
    dv_list(ws, 'C12:E12', '物料种类表')
    win = _win(p['start'], p['end'])
    neg = _neg_fmt(wb)
    for r in range(14, 214):
        b = f'$B{r}'
        f = {'A': f'=IF({b}="","",ROW()-13)',
             'B': f'=IFERROR(INDEX({AU}!$DP$4:$DP$203,ROW()-13),"")'}
        for col in mats:
            f[col] = (f'=IF(OR({b}="",{_hdr_blank(col + "$12")}),"",SUMIFS({au("DA")},{bz("D")},{b},'
                      f'{bz("E")},{col}$12,{win}))')
        f['F'] = f'=IF({b}="","",SUM($C{r}:$E{r}))'
        f['G'] = f'=IF({b}="","",SUMIFS({au("DB")},{bz("D")},{b},{au("CZ")},"发泡网",{win}))'
        f['H'] = f'=IF({b}="","",N($F{r})-N($G{r}))'
        f['I'] = f'=IF({b}="","",ROUND(SUMIFS({au("DC")},{bz("D")},{b},{au("CZ")},"发泡网",{win}),2))'
        f['J'] = f'=IF({b}="","",ROUND(SUMIFS({au("DD")},{bz("D")},{b},{au("CZ")},"发泡网",{win}),2))'
        f['K'] = f'=IF({b}="","",ROUND(N($I{r})-N($J{r}),2))'
        f['L'] = None
        for col, v in f.items():
            fmt = MONEY if col in 'IJK' else (neg if col == 'G' else (INT if col == 'A' else QTY))
            cell(ws, f'{col}{r}', v, F_IN if col == 'L' else F_AUTO, FL_IN if col == 'L' else FL_AUTO,
                 AL if col in 'BL' else AC, fmt=fmt)
    for t in (13, 214):
        cell(ws, f'A{t}', None, F_TOT, FL_TOT, AC)
        cell(ws, f'B{t}', '合计', F_TOT, FL_TOT, AC)
        for col in 'CDEFGHIJK':
            cell(ws, f'{col}{t}', f'=ROUND(SUM({col}14:{col}213),2)', F_TOTN, FL_TOT, AC,
                 fmt=MONEY if col in 'IJK' else (neg if col == 'G' else QTY))
        cell(ws, f'L{t}', None, F_TOT, FL_TOT, AC)
    widths(ws, {'A': 10, 'B': 13, 'C': 11, 'D': 11, 'E': 11, 'F': 11, 'G': 11, 'H': 12, 'I': 12, 'J': 12,
                'K': 12, 'L': 13})
    ws.freeze_panes = 'C14'


def build_qtw(wb, idx):
    ws = _new_sheet(wb, QTW, idx)
    title(ws, '其 他 包 装 物 料 汇 总（发泡网、筐子以外 · 库存成本＋按客户 · 全自动）', 'A', 'R')
    p = period(ws, row=2, col='A', open_ended=True, year=None)
    note(ws, '★ 全自动。物料范围＝【包装物料出入库明细】里不是发泡网（见【发泡网销售汇总】第 12 行蓝格）、也不是筐子/托盘的物料。'
             '上块按物料列库存与成本，加权单价含期初（期初那几种物料不再是 0 成本）；下块每个客户一行，C～J 按第 29 行表头'
             '（蓝格，可改）统计数量，K「其他」＝没列出来的。销售金额＝领用＋销售−退回，已收＝已收讫。第 2 行可按区间统计，都空＝不限。',
         'A', 'R', 3, 48)
    section(ws, '一、库存与成本（按物料）', 'A', 'L', 4)
    _stock_block(ws, 6, 25, lambda r: f'=IFERROR(INDEX({AU}!$EK$4:$EK$63,ROW()-5),"")', p)
    section(ws, '二、按客户（领用出库＋销售出库，客户退回冲减）', 'A', 'R', 28)
    mats = [L(i) for i in range(3, 11)]       # C～J
    names = ['香梨纸', '无纺布', '气泡垫', '纸垫板', '胶带', '冷煤/氟利昂', BAK, BAK]
    header(ws, 29, 'A', ['序号', '客户'] + names + ['其他', '合计数量', '客户退回\n(显示为负)', '净数量',
                                                  '销售金额', '已收金额', '未收金额', '备注'])
    for col in mats:
        cell(ws, f'{col}29', ws[f'{col}29'].value, F_IN, FL_IN, AC)
    _red_header_cf(ws, 'C29:J29', f'{JC}!$A$4:$A$103')
    dv_list(ws, 'C29:J29', '物料种类表')
    win = _win(p['start'], p['end'])
    neg = _neg_fmt(wb)
    for r in range(31, 231):
        b = f'$B{r}'
        f = {'A': f'=IF({b}="","",ROW()-30)',
             'B': f'=IFERROR(INDEX({AU}!$DV$4:$DV$203,ROW()-30),"")'}
        for col in mats:
            f[col] = (f'=IF(OR({b}="",{_hdr_blank(col + "$29")}),"",SUMIFS({au("DA")},{bz("D")},{b},'
                      f'{bz("E")},{col}$29,{win}))')
        f['K'] = f'=IF({b}="","",SUMIFS({au("DA")},{bz("D")},{b},{au("CZ")},"其他",{win})-SUM($C{r}:$J{r}))'
        f['L'] = f'=IF({b}="","",SUM($C{r}:$K{r}))'
        f['M'] = f'=IF({b}="","",SUMIFS({au("DB")},{bz("D")},{b},{au("CZ")},"其他",{win}))'
        f['N'] = f'=IF({b}="","",N($L{r})-N($M{r}))'
        f['O'] = f'=IF({b}="","",ROUND(SUMIFS({au("DC")},{bz("D")},{b},{au("CZ")},"其他",{win}),2))'
        f['P'] = f'=IF({b}="","",ROUND(SUMIFS({au("DD")},{bz("D")},{b},{au("CZ")},"其他",{win}),2))'
        f['Q'] = f'=IF({b}="","",ROUND(N($O{r})-N($P{r}),2))'
        f['R'] = None
        for col, v in f.items():
            fmt = MONEY if col in 'OPQ' else (neg if col == 'M' else (INT if col == 'A' else QTY))
            cell(ws, f'{col}{r}', v, F_IN if col == 'R' else F_AUTO, FL_IN if col == 'R' else FL_AUTO,
                 AL if col in 'BR' else AC, fmt=fmt)
    for t in (30, 231):
        cell(ws, f'A{t}', None, F_TOT, FL_TOT, AC)
        cell(ws, f'B{t}', '合计', F_TOT, FL_TOT, AC)
        for col in 'CDEFGHIJKLMNOPQ':
            cell(ws, f'{col}{t}', f'=ROUND(SUM({col}31:{col}230),2)', F_TOTN, FL_TOT, AC,
                 fmt=MONEY if col in 'OPQ' else (neg if col == 'M' else QTY))
        cell(ws, f'R{t}', None, F_TOT, FL_TOT, AC)
    widths(ws, {'A': 10, 'B': 13, **{L(i): 11 for i in range(3, 15)}, 'O': 12, 'P': 12, 'Q': 12, 'R': 13})
    ws.freeze_panes = 'C31'


def build_cgh(wb, idx):
    ws = _new_sheet(wb, CGH, idx, tab='10243E')
    title(ws, '包 装 物 料 采 购 汇 总（按供应商 · 取自【包装物料出入库明细】业务类型＝采购入库）', 'A', 'R')
    p = period(ws, row=2, col='I', open_ended=True, year=None)
    note(ws, '★ 全自动。C～L 按第 3 行表头（蓝格，可改，新品直接改「（备用）」）统计采购数量，M「其他」＝没列出来的物料。'
             '已付金额＝收付款状态为「已付款」的部分（真实付款在《03》资金日记账）。', 'A', 'H', 2, 48)
    mats = [L(i) for i in range(3, 13)]       # C～L
    names = ['香梨网', '苹果网', '无纺布', '香梨纸', '气泡垫', '纸垫板', '冷煤/氟利昂', BAK, BAK, BAK]
    header(ws, 3, 'A', ['序号', '供应商'] + names + ['其他', '合计数量', '采购金额', '已付金额', '未付金额', '备注'], 36)
    for col in mats:
        cell(ws, f'{col}3', ws[f'{col}3'].value, F_IN, FL_IN, AC)
    _red_header_cf(ws, 'C3:L3', f'{JC}!$A$4:$A$103')
    dv_list(ws, 'C3:L3', '物料种类表')
    win = _win(p['start'], p['end'])
    for r in range(5, 205):
        b = f'$B{r}'
        cg_ = lambda col, extra='': f'SUMIFS({bz(col)},{bz("D")},{b},{bz("H")},"采购入库"{extra},{win})'
        f = {'A': f'=IF({b}="","",ROW()-4)',
             'B': f'=IFERROR(INDEX({AU}!$EB$4:$EB$203,ROW()-4),"")'}
        for col in mats:
            f[col] = f'=IF(OR({b}="",{_hdr_blank(col + "$3")}),"",{cg_("I", "," + bz("E") + "," + col + "$3")})'
        f['M'] = f'=IF({b}="","",{cg_("I")}-SUM($C{r}:$L{r}))'
        f['N'] = f'=IF({b}="","",SUM($C{r}:$M{r}))'
        f['O'] = f'=IF({b}="","",ROUND({cg_("N")},2))'
        paid = ',' + bz('P') + ',"已付款"'
        f['P'] = f'=IF({b}="","",ROUND({cg_("N", paid)},2))'
        f['Q'] = f'=IF({b}="","",ROUND(N($O{r})-N($P{r}),2))'
        f['R'] = None
        for col, v in f.items():
            fmt = MONEY if col in 'OPQ' else (INT if col == 'A' else QTY)
            cell(ws, f'{col}{r}', v, F_IN if col == 'R' else F_AUTO, FL_IN if col == 'R' else FL_AUTO,
                 AL if col in 'BR' else AC, fmt=fmt)
    for t in (4, 205):
        cell(ws, f'A{t}', None, F_TOT, FL_TOT, AC)
        cell(ws, f'B{t}', '合计', F_TOT, FL_TOT, AC)
        for col in 'CDEFGHIJKLMNOPQ':
            cell(ws, f'{col}{t}', f'=ROUND(SUM({col}5:{col}204),2)', F_TOTN, FL_TOT, AC,
                 fmt=MONEY if col in 'OPQ' else QTY)
        cell(ws, f'R{t}', None, F_TOT, FL_TOT, AC)
    widths(ws, {'A': 6, 'B': 14, **{L(i): 10 for i in range(3, 15)}, 'O': 12, 'P': 12, 'Q': 12, 'R': 14})
    ws.freeze_panes = 'C5'


def build_cost_iface(wb, idx):
    """新表【物料成本接口】（K7）：B2 起始月；行 4～27＝24 个月；月末一次加权平均；隐藏网格 K～Z 行 4～1443"""
    ws = _new_sheet(wb, CBJ, idx, tab='C55A11')
    title(ws, '物 料 成 本 接 口（全自动 · 供《03 财务账套》按月结转物料销售成本，请勿改动结构）', 'A', 'G')
    cell(ws, 'A2', '起始月', F_LBL, FL_LBL, AC)
    cell(ws, 'B2', datetime.datetime(2026, 8, 1), F_IN, FL_IN, AC, fmt=MONTH)
    ws.merge_cells('C2:G2')
    cell(ws, 'C2', '← 蓝格可改（填每月 1 日，默认 2026-08-01，与《03》对接源_02物料 同一个起点）。成本按「月末一次加权平均」：'
                   '每种物料 当月单价＝(月初结存金额＋本月采购金额)÷(月初结存数量＋本月采购数量)，销售成本＝本月净销售数量'
                   '(领用＋销售−客户退回)×当月单价；建账期初按其金额进第一个月的月初结存。「哪些算发泡网」看【发泡网销售汇总】第 12 行。',
         F_NOTE, FL_NONE, AL, border=NOB)
    ws.row_dimensions[2].height = 54
    header(ws, 3, 'A', ['月份', '发泡网\n销售数量', '发泡网\n销售成本', '其他包装物料\n销售数量', '其他包装物料\n销售成本',
                         '销售成本合计', '备注'], 36)
    G0, G1 = 4, 4 + 60 * 24 - 1             # 网格 4～1443：60 种物料 × 24 个月
    rng = lambda col: f'${col}${G0}:${col}${G1}'
    for i in range(24):
        r = 4 + i
        cell(ws, f'A{r}', f'=DATE(YEAR($B$2),MONTH($B$2)+{i},1)', F_AUTOB, FL_AUTO, AC, fmt=MONTH)
        for col, src, cat in (('B', 'U', '发泡网'), ('C', 'V', '发泡网'), ('D', 'U', '其他'), ('E', 'V', '其他')):
            cell(ws, f'{col}{r}', f'=ROUND(SUMIFS({rng(src)},{rng("O")},"{cat}",{rng("M")},$A{r}),2)',
                 F_AUTO, FL_AUTO, AC, fmt=MONEY if col in 'CE' else QTY)
        cell(ws, f'F{r}', f'=ROUND(N($C{r})+N($E{r}),2)', F_AUTOB, FL_AUTO, AC, fmt=MONEY)
        cell(ws, f'G{r}', None, F_IN, FL_IN, AL)
    cell(ws, 'A28', '合计', F_TOT, FL_TOT, AC)
    for col in 'BCDEF':
        cell(ws, f'{col}28', f'=ROUND(SUM({col}4:{col}27),2)', F_TOTN, FL_TOT, AC, fmt=MONEY if col in 'CEF' else QTY)
    cell(ws, 'G28', None, F_TOT, FL_TOT, AC)
    # 隐藏网格
    gh = ['物料', '月序', '月初日', '下月初', '类别', '月初数量', '月初金额', '本月采购数量', '本月采购金额', '当月单价',
          '净销售数量', '销售成本', '其他减少数量', '其他减少成本', '月末数量', '月末金额']
    for i, t in enumerate(gh):
        ws.cell(3, 11 + i, t).font = F_HELP
    H = bz('H')
    c_qc, c_nqc, c_cg = ',%s,"期初库存"' % H, ',%s,"<>期初库存"' % H, ',%s,"采购入库"' % H
    for r in range(G0, G1 + 1):
        k = f'$K{r}'
        mw = ',' + bz('B') + ',">="&$M' + str(r) + ',' + bz('B') + ',"<"&$N' + str(r)
        pre = ',' + bz('B') + ',"<"&$M' + str(r)

        def s(col, *extra):
            return 'SUMIFS(%s,%s,%s%s)' % (bz(col), bz('E'), k, ''.join(extra))
        v = {
            'K': f'=IFERROR(INDEX({AU}!$G$4:$G$63,INT((ROW()-{G0})/24)+1),"")',
            'L': f'=MOD(ROW()-{G0},24)',
            'M': f'=IF({k}="","",DATE(YEAR($B$2),MONTH($B$2)+$L{r},1))',
            'N': f'=IF({k}="","",DATE(YEAR($B$2),MONTH($B$2)+$L{r}+1,1))',
            'O': '=' + _cat(k),
            'P': (f'=IF({k}="","",IF($L{r}=0,{s("I", c_qc)}+{s("I", c_nqc, pre)}-{s("J", pre)}+{s("K", pre)},'
                  f'N($Y{r - 1})))'),
            'Q': (f'=IF({k}="","",IF($L{r}=0,IF(N($P{r})=0,0,ROUND(N($P{r})*IFERROR(({s("N", c_qc)}+{s("N", c_cg, pre)})/'
                  f'({s("I", c_qc)}+{s("I", c_cg, pre)}),0),2)),N($Z{r - 1})))'),
            'R': f'=IF({k}="","",{s("I", c_cg, mw)})',
            'S': f'=IF({k}="","",{s("N", c_cg, mw)})',
            'T': f'=IF({k}="","",IF(N($P{r})+N($R{r})>0,(N($Q{r})+N($S{r}))/(N($P{r})+N($R{r})),0))',
            'U': f'=IF({k}="","",SUMIFS({au("DA")},{bz("E")},{k}{mw})-SUMIFS({au("DB")},{bz("E")},{k}{mw}))',
            'V': f'=IF({k}="","",ROUND(N($U{r})*N($T{r}),2))',
            'W': f'=IF({k}="","",{s("J", mw)}-SUMIFS({au("DA")},{bz("E")},{k}{mw}))',
            'X': f'=IF({k}="","",ROUND(N($W{r})*N($T{r}),2))',
            'Y': f'=IF({k}="","",N($P{r})+N($R{r})-N($U{r})-N($W{r}))',
            'Z': f'=IF({k}="","",IF(N($Y{r})=0,0,ROUND(N($Q{r})+N($S{r})-N($V{r})-N($X{r}),2)))',
        }
        for col, f in v.items():
            c = ws[f'{col}{r}']
            c.value = f
            c.font = F_HELP
            if col in 'MN':
                c.number_format = DATE
    widths(ws, {'A': 12, 'B': 12, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 24})
    hide(ws, *[L(i) for i in range(11, 27)])
    ws.freeze_panes = 'B4'


# ════════════════════════════════════════════════════════════════════════════
#  八、主页、使用说明
# ════════════════════════════════════════════════════════════════════════════
HOME = [
    ('一、日常录入（每天要填的表）', [
        (BZ, '物料进出与销售；计价数量和合计金额自动算', '手工录入'),
        (TZ, '公司筐的采购、发出、客户退回、卖断；剩余库存按类型自动滚动；含收退押金；O 列装筐费自动算', '手工录入'),
        (ZB, '客户自己带来的筐子，表外备查', '手工录入'),
        (CG, '次果筐与托盘合并在这一张；含押金与退款情况', '手工录入'),
        ('实盘库存盘点', '只在右半边录实盘数量，左半边账面数与盘盈盘亏全自动', '半自动'),
    ]),
    ('二、库存与往来（全部自动，不用填）', [
        ('物料库存结余', '按物料看期初/采购/退回/领用/销售与剩余库存、库存金额（加权单价含期初）', '全自动'),
        ('筐子库存汇总', '公司每种筐子还剩多少、客户手上多少、公司筐总量', '全自动'),
        ('筐子往来汇总', '按客户/供应商看发出、退回（含次果筐退回）、采购，已含01册带出带回的筐数', '全自动'),
        ('筐子采购汇总', '按供应商看采购筐子的数量、金额、已付未付', '全自动'),
        ('筐子销售汇总', '卖断给客户的筐子数量、金额、已收未收（类型看表头蓝格）', '全自动'),
        ('自备筐汇总', '客户自备筐进出与结存，联动01册成品出库带走的自备筐', '全自动'),
        ('次果筐汇总', '次果筐/托盘按客户的出库、退回、未回与押金', '全自动'),
        ('押金汇总', '每家客户收了多少、退了多少、还欠客户多少押金', '全自动'),
        ('客户领用物料汇总', '按客户看各物料领用量（表头蓝格可改）、退回、净领用与应收已收', '全自动'),
        (FPW, '【新】香梨网/苹果网 库存、加权成本、按客户的销售与收款', '全自动'),
        (QTW, '【新】发泡网、筐子以外的包装物料：库存成本＋按客户销售', '全自动'),
        (CGH, '【新】按供应商看包装物料采购数量、金额、已付未付', '全自动'),
        ('收入结算汇总', '筐子销售、物料销售两项收入 ＋ 押金专区 ＋ 期末存货', '全自动'),
    ]),
    ('三、对接与码表', [
        (DY, '跨文件取《01 水果进销存台账》的筐子进出数，断链自动切备份', '全自动'),
        (JK, '四张明细逐笔规范化（3,000 行），给《04》对账单用', '全自动'),
        (ZKF, '【新】装筐费逐笔清单，给《03》挂应付', '全自动'),
        (CBJ, '【新】24 个月发泡网/包装物料销售成本（月末加权），给《03》结转', '全自动'),
        ('财务取数接口', '按月汇总的旧接口（《03》已不再引用，留作备查）', '全自动'),
        (JC, '物料、筐子类型、客户、供应商等下拉选项；U～V 列装筐费单价', '码表'),
        ('字段对照', '原有 Excel 字段 → 本册字段的对照表，存档备查', '说明'),
    ]),
]


def home(wb):
    ws = wb['主页']
    st = {k: copy.copy(ws[k]._style) for k in ('A5', 'A6', 'B6', 'C6', 'D6', 'A7', 'B7', 'C7', 'D7', 'D11', 'D15', 'D29', 'D30', 'A32')}
    h_sec, h_row = ws.row_dimensions[5].height, ws.row_dimensions[7].height
    legend = ws['A32'].value
    for m in [str(x) for x in ws.merged_cells.ranges]:
        if int(re.search(r'\d+', m).group()) >= 5:
            ws.unmerge_cells(m)
    for (r, c), x in list(ws._cells.items()):
        if r >= 5:
            del ws._cells[(r, c)]
    for k in [k for k in ws.row_dimensions if k >= 5]:
        del ws.row_dimensions[k]
    attr_st = {'手工录入': 'D7', '半自动': 'D11', '全自动': 'D15', '码表': 'D29', '说明': 'D30'}
    r, n = 5, 0
    for sec, items in HOME:
        ws.merge_cells(f'A{r}:D{r}')
        ws[f'A{r}'] = sec
        ws[f'A{r}']._style = copy.copy(st['A5'])
        ws.row_dimensions[r].height = h_sec
        r += 1
        for col, t in zip('ABCD', ['序号', '子表名称', '用途说明', '属性']):
            ws[f'{col}{r}'] = t
            ws[f'{col}{r}']._style = copy.copy(st[f'{col}6'])
        r += 1
        for name, desc, attr in items:
            n += 1
            ws[f'A{r}'] = n
            ws[f'B{r}'] = name
            ws[f'B{r}'].hyperlink = Hyperlink(ref=f'B{r}', location=f"'{name}'!A1", display=name)
            ws[f'C{r}'] = desc
            ws[f'D{r}'] = attr
            for col, k in (('A', 'A7'), ('B', 'B7'), ('C', 'C7'), ('D', attr_st[attr])):
                ws[f'{col}{r}']._style = copy.copy(st[k])
            ws.row_dimensions[r].height = h_row
            r += 1
        r += 1
    ws.merge_cells(f'A{r}:D{r}')
    ws[f'A{r}'] = legend
    ws[f'A{r}']._style = copy.copy(st['A32'])
    ws.print_area = f'A1:D{r}'


def manual(wb):
    ws = wb['使用说明']
    last = max(r for (r, c), x in ws._cells.items() if x.value not in (None, ''))
    st_h, st_t = copy.copy(ws['A84']._style), copy.copy(ws['A85']._style)
    lines = [
        '★ 本 次 更 新（2026-10 · 第九版）',
        '1、所有汇总表第 4 行加了「合计」（数据从第 5 行起，底部合计照旧）；新表：发泡网销售汇总、其他包装物料汇总、包装物料采购汇总、装筐费接口、物料成本接口（主页有入口）。',
        '2、【客户领用物料汇总】【筐子销售汇总】【包装物料采购汇总】等的物料/筐子列按第 3 行（或表头行）的蓝色表头统计：想看哪种就把表头改成哪种，表头在清单里找不到会标红；没列出来的进「其他」。客户领用物料的应收＝领用＋销售−退回，跟《03》一致。',
        '3、【周转筐出入库明细】在「装卸」后面加了 O「装筐费」＝（入库数量＋退回数量）×单价，单价在【基础资料】U～V 列填（留空＝不计费），期初和 自卸/自装/自提/无 不计（U18 起可改）。原 O～W 往右挪了一列。',
        '4、【筐子往来汇总】的退回含【次果筐+托盘明细】里「退回入库」的筐。',
        '5、扩容：周转筐出入库明细可录到第 3003 行，对账明细接口 3,000 行。所有明细表请在最后一行下面接着录，不要在中间插行。',
    ]
    for i, t in enumerate(lines):
        r = last + 2 + i
        c = ws[f'A{r}']
        c.value = t
        c._style = copy.copy(st_h if i == 0 else st_t)
        ws.merge_cells(f'A{r}:H{r}')
        ws.row_dimensions[r].height = 19.5 if i == 0 else (33.75 if len(t) < 90 else 48)


# ════════════════════════════════════════════════════════════════════════════
#  九、入口
# ════════════════════════════════════════════════════════════════════════════
def apply(wb):
    structure(wb)                 # 插列、扩容、池子、对账明细接口、顶部合计（只动位置，数不变）
    fee_param(wb)
    tz_fee(wb)
    new_pools(wb)
    fixes(wb)
    collect_material(wb)
    basket_sales(wb)
    basket_flow(wb)
    i = wb.sheetnames.index('实盘库存盘点')
    build_fpw(wb, i)
    build_qtw(wb, i + 1)
    build_cgh(wb, i + 2)
    j = wb.sheetnames.index('财务取数接口')
    build_fee_iface(wb, j)
    build_cost_iface(wb, j + 1)
    home(wb)
    manual(wb)


def post(path):
    """本模块不需要 XML 层后处理"""
    return None
