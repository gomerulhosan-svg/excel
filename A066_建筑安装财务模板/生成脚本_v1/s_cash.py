# -*- coding: utf-8 -*-
"""【资金流水】所有账户一张总表（九江银行、农业银行、工商银行、现金、老板/负责人的个人户、承兑汇票）。
   跟原来一样录：日期、账户、摘要、收入、支出；紧跟着自动认出 项目 / 收支类别 / 单位或人，认错了在右边黄格改。"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *


def hit(text, sorted_rng):
    """摘要里对上的第一个关键词在排好的关键词表（长的在前）里是第几个，没对上＝#N/A。
       只用 MATCH(1, INDEX(ISNUMBER(SEARCH(区域, 文字))*(区域<>""), 0), 0)：Excel、WPS、LibreOffice 都按数组算"""
    return f'MATCH(1,INDEX(ISNUMBER(SEARCH({sorted_rng},{text}))*({sorted_rng}<>""),0),0)'


def first_hit(text, sorted_rng):
    """同上，没对上＝0（用的地方先判断 0，别让 INDEX(区域,0) 取到整列）"""
    return f'IFERROR({hit(text, sorted_rng)},0)'


FA_NAMES = rng(SH_BASE, FA_NAME, FA_R0, FA_R1)
FA_DATES = rng(SH_BASE, FA_DATE, FA_R0, FA_R1)


def build_cash(wb, ctx):
    ws = wb.create_sheet(SH_CASH)
    widths(ws, {'A': 6, 'B': 11, 'C': 11, 'D': 34, 'E': 14, 'F': 14, 'G': 13, 'H': 12, 'I': 11, 'J': 11, 'K': 30, 'L': 12,
                'M': 11, 'N': 11, 'O': 11, 'P': 14})
    title(ws, '资 金 流 水（所有账户一张表 · 每天录 · 自动认项目、类别、单位）', J_NOTE, C_CASH,
          '💡 左边照原来录：日期、账户、摘要、收入、支出（收款记「收入」，付款记「支出」；银行退回来的钱在支出栏记负数也行，自动当冲回）。'
          '中间灰色是按摘要自动认出的 项目 / 收支类别 / 单位或人，看一眼对不对；不对或认不出，在右边黄格「改项目、改类别、改单位」选（改的优先）。'
          '几条规矩：①「代X付」（代老叶付佩力材料款）：单位是 X（冲 X 的账），真正收钱的写备注；'
          '② 一笔钱是几个项目的（隆世光伏/共青办公楼/德祥工地工程款），拆成几行分别记；'
          '③ 自己账户之间转钱（银行↔银行、银行↔老板个人户：借款、还款、报销、备用金）类别是「账户互转」，只记一边，对方账户自动认或在 O 列选；'
          '④ 老板、负责人替公司花的钱，记在他的个人户里（账户选他的个人户）。G 列是这个账户记到这一笔为止的余额。'
          '⑤ 记错的行清空内容就行，不要删整行、也不要插行（别的表按行取数）。')
    put(ws, 'A3', '笔数', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B3', f'=COUNT({jr(J_DATE)})', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'C3', '收入合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'D3', f'=SUMIFS({jr(J_NET)},{jr(J_IO)},"收",{jr(J_CT)},"<>账户互转")', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    put(ws, 'E3', '支出合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'F3', f'=-SUMIFS({jr(J_NET)},{jr(J_IO)},"支",{jr(J_CT)},"<>账户互转")', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    put(ws, 'G3', '待处理', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'H3', f'=COUNTIF({jr(J_CHK)},"✗*")&" 笔"', F_KPI_V, fill('FFD9E1F2'), align=AC)
    put(ws, 'I3', '提醒', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'J3', f'=COUNTIF({jr(J_CHK)},"⚠*")&" 笔"', F_KPI_V, fill('FFD9E1F2'), align=AC)
    put(ws, 'K3', f'="收支不含账户互转；容量 "&TEXT(B3/{J_CAP},"0%")', F_NOTE, fill('FFD9E1F2'), align=AL)
    for c1, c2, t, col in (('A', 'F', '① 照原来录', 'FF305496'), ('G', 'K', '② 自动认的（看对不对）', 'FF7F7F7F'),
                           ('L', 'P', '③ 不对才改（改的优先）', 'FFBF8F00')):
        section(ws, 4, c1, c2, t, col)
    heads = [(J_SEQ, '序号'), (J_DATE, '日期'), (J_ACC, '账户'), (J_MEMO, '摘要'), (J_IN, '收入'), (J_OUT, '支出'),
             (J_BAL, '结存\n（本账户）'), (J_PJ, '项目'), (J_CT, '收支类别'), (J_UN, '单位 / 人'), (J_CHK, '校验'),
             (J_MPJ, '改项目'), (J_MCT, '改类别'), (J_MUN, '改单位/人'), (J_TOACC, '对方账户\n（互转）'), (J_NOTE, '备注')]
    for col, t in heads:
        c = 'FF305496' if CI(col) <= CI(J_OUT) else ('FF7F7F7F' if CI(col) <= CI(J_CHK) else 'FFBF8F00')
        put(ws, f'{col}{J_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[J_HDR].height = 34
    hidden = [(J_NET, '净额'), (J_YM, '年月'), (J_APJ, '认项目'), (J_ACT, '认类别'), (J_AUN, '认单位'), (J_ALLOC, '分摊归类'),
              (J_LINE, '报表项目'), (J_HANG, '有应付'), (J_UTYPE, '单位类型'), (J_WKIND, '走考勤的人'), (J_DAI, '代X付'),
              (J_ATO, '认对方账户'), (J_TO, '对方账户'), (J_ATYPE, '账户类型'), (J_TOTYPE, '对方类型'), (J_KWC, '关键词类别'),
              (J_IO, '收支'), (J_PKW, '项目命中词'), (J_PKN, '别的项目'), (J_PKH, '项目码'), (J_UMH, '单位码'), (J_UDH, '代付单位码'), (J_DN, '日期数'),
              (J_RK1, '收款排号'), (J_RK2, '付款排号'), (J_RK3, '押金排号')]
    for col, t in hidden:
        ws[f'{col}{J_HDR}'] = t
        ws[f'{col}{J_HDR}'].font = F_HELP
    AP_UNITS, OAP_UNITS = apr(AP_UNIT), oapr(OAP_UNIT)
    for r in range(J_R0, J_R1 + 1):
        g = lambda c: f'{c}{r}'
        memo, un, ut, ct = g(J_MEMO), g(J_UN), g(J_UTYPE), g(J_CT)
        empty = f'AND({g(J_DATE)}="",{memo}="",{g(J_IN)}="",{g(J_OUT)}="")'
        f = {}
        f[J_SEQ] = f'=IF({empty},"",ROW()-{J_HDR})'
        f[J_NET] = f'=ROUND(N({g(J_IN)})-N({g(J_OUT)}),2)'
        f[J_YM] = f'=IF(ISNUMBER({g(J_DATE)}),YEAR({g(J_DATE)})*100+MONTH({g(J_DATE)}),"")'
        f[J_DN] = f'=IF(ISNUMBER({g(J_DATE)}),INT({g(J_DATE)}),0)'
        f[J_BAL] = (f'=IF({g(J_ACC)}="","",ROUND(IFERROR(INDEX({AC_OPENS},MATCH({g(J_ACC)},{AC_NAMES},0)),0)'
                    f'+SUMIFS(${J_NET}${J_R0}:{g(J_NET)},${J_ACC}${J_R0}:{g(J_ACC)},{g(J_ACC)})'
                    f'-SUMIFS(${J_NET}${J_R0}:{g(J_NET)},${J_TO}${J_R0}:{g(J_TO)},{g(J_ACC)},${J_LINE}${J_R0}:{g(J_LINE)},"账户互转"),2))')
        f[J_PKH] = f'=IF({memo}="",0,{first_hit(memo, AX_PS_R)})'
        f[J_PKW] = f'=IF(N({g(J_PKH)})=0,"",INDEX({AX_PS_R},{g(J_PKH)}))'
        f[J_APJ] = f'=IF(N({g(J_PKH)})=0,"",INDEX({AX_PN_R},{g(J_PKH)}))'
        f[J_PKN] = (f'=IF(OR({g(J_PKW)}="",{g(J_MPJ)}<>""),0,SUMPRODUCT(ISNUMBER(SEARCH({AX_PS_R},{memo}))*({AX_PS_R}<>"")'
                    f'*(1-ISNUMBER(SEARCH({AX_PS_R},{g(J_PKW)})))*({AX_PN_R}<>{g(J_APJ)})))')
        f[J_DAI] = (f'=IF(ISERROR(FIND("代",{memo})),"",IFERROR(MID({memo},FIND("代",{memo})+1,'
                    f'FIND("付",{memo},FIND("代",{memo}))-FIND("代",{memo})-1),""))')
        dai = g(J_DAI)
        f[J_UMH] = f'=IF({memo}="",0,{first_hit(memo, AX_US_R)})'
        f[J_UDH] = f'=IF({dai}="",0,{first_hit(dai, AX_US_R)})'
        um, ud = (f'IF(N({g(h)})=0,"",INDEX({AX_UN_R},{g(h)}))' for h in (J_UMH, J_UDH))
        f[J_AUN] = f'=IF({memo}="","",IF(N({g(J_UDH)})>0,{ud},{um}))'
        f[J_PJ] = f'=IF({g(J_MPJ)}<>"",{g(J_MPJ)},{g(J_APJ)})'
        f[J_UN] = f'=IF({g(J_MUN)}<>"",{g(J_MUN)},{g(J_AUN)})'
        f[J_UTYPE] = f'=IF({un}="","",IFERROR(INDEX({UN_TYPES_R},MATCH({un},{UN_NAMES},0))&"",""))'
        f[J_WKIND] = f'=IF(OR({ut}="管理人员",{ut}="工人"),1,0)'
        f[J_HANG] = f'=IF({un}="",0,IF(COUNTIF({AP_UNITS},{un})+COUNTIF({OAP_UNITS},{un})>0,1,0))'
        dirn = f'IF({g(J_OUT)}<>"","支","收")'
        nkw = KW_R1 - KW_R0 + 1
        kh = lambda col: hit(memo, rng(SH_AUX, col, 1, nkw))
        kw = f'IFERROR(INDEX({KW_CATS},IF({dirn}="收",{kh(AX_KWI)},{kh(AX_KWO)}))&"","")'
        f[J_KWC] = f'=IF(OR({memo}="",{g(J_NET)}=0),"",{kw})'
        k = g(J_KWC)
        dflt = (f'IF({g(J_WKIND)}=1,"工资",IF({ut}="临时工","临时工工资",IF({ut}="甲方/总包","工程款",'
                f'IF(AND(OR({ut}="材料供应商",{ut}="分包",{ut}="机械运输"),{g(J_HANG)}=1),"付应付款",""))))')
        hangable = (f'AND({g(J_HANG)}=1,OR({ut}="材料供应商",{ut}="分包",{ut}="机械运输",{ut}="其他"),'
                    f'OR({k}="材料费",{k}="分包费",{k}="机械运输费",{k}="其他直接费",{k}="工资",{k}="临时工工资"))')
        f[J_ACT] = (f'=IF(OR({memo}="",{g(J_NET)}=0),"",IF(AND({ut}="临时工",{g(J_NET)}<0),"临时工工资",'
                    f'IF({k}="",{dflt},IF({hangable},"付应付款",{k}))))')
        f[J_CT] = f'=IF({g(J_MCT)}<>"",{g(J_MCT)},{g(J_ACT)})'
        f[J_ALLOC] = f'=IF({ct}="","",IFERROR(INDEX({CT_ALLOCS},MATCH({ct},{CT_NAMES},0))&"","？"))'
        cdir = f'IFERROR(INDEX({CT_DIRS},MATCH({ct},{CT_NAMES},0))&"","")'
        f[J_IO] = f'=IF({g(J_NET)}=0,"",IF(OR({cdir}="收",{cdir}="支"),{cdir},IF({g(J_NET)}>0,"收","支")))'
        ah = first_hit(memo, AX_AS_R)
        am = f'IF({ah}=0,"",INDEX({AX_AN_R},{ah}))'
        f[J_ATO] = f'=IF({ct}<>"账户互转","",IF({un}<>"",IFERROR(INDEX({AX_AKA_R},MATCH({un},{AX_AKW_R},0)),{am}),{am}))'
        f[J_TO] = f'=IF({g(J_TOACC)}<>"",{g(J_TOACC)},{g(J_ATO)})'
        f[J_ATYPE] = f'=IF({g(J_ACC)}="","",IFERROR(INDEX({AC_TYPES_R},MATCH({g(J_ACC)},{AC_NAMES},0))&"",""))'
        f[J_TOTYPE] = f'=IF({g(J_TO)}="","",IFERROR(INDEX({AC_TYPES_R},MATCH({g(J_TO)},{AC_NAMES},0))&"",""))'
        apl = f'IF({ut}="材料供应商","应付材料款",IF({ut}="分包","应付分包款",IF({ut}="机械运输","应付机械运输费","应付其他款")))'
        lk = f'IFERROR(INDEX({CT_LINES},MATCH({ct},{CT_NAMES},0))&"","")'
        f[J_LINE] = (f'=IF({g(J_NET)}=0,"",IF(OR({ct}="",{lk}=""),"待分类",IF(OR({ct}="付应付款",{lk}="应付款"),{apl},{lk})))')
        # 「最后一次」排号：这一笔是这个项目/单位/人的第几次（按日期，同一天的算同一号）
        dn = g(J_DN)
        upto = f'{jr(J_DN)},">0",{jr(J_DN)},"<="&{dn}'
        f[J_RK1] = (f'=IF(AND({g(J_LINE)}="应收账款",{g(J_NET)}>0,{dn}>0,{g(J_PJ)}<>""),{g(J_PJ)}&"|"&'
                    f'COUNTIFS({jr(J_PJ)},{g(J_PJ)},{jr(J_LINE)},"应收账款",{jr(J_NET)},">0",{upto}),"")')
        f[J_RK2] = (f'=IF(AND({un}<>"",{g(J_NET)}<0,{dn}>0),{un}&"|"&'
                    f'COUNTIFS({jr(J_UN)},{un},{jr(J_NET)},"<0",{upto}),"")')
        # 押金按行取项目：同一天几笔再按上下排（取最下面那笔）
        f[J_RK3] = (f'=IF(AND({g(J_LINE)}="其他应收款",{dn}>0,{un}<>""),{un}&"|"&'
                    f'(COUNTIFS({jr(J_UN)},{un},{jr(J_LINE)},"其他应收款",{jr(J_DN)},">0",{jr(J_DN)},"<"&{dn})'
                    f'+COUNTIFS(${J_UN}${J_R0}:{un},{un},${J_LINE}${J_R0}:{g(J_LINE)},"其他应收款",${J_DN}${J_R0}:{dn},{dn})),"")')
        need = f'IFERROR(INDEX({CT_NEEDS},MATCH({ct},{CT_NAMES},0))&"","")'
        ln = g(J_LINE)
        direct = f'OR({ln}="材料费",{ln}="人工费",{ln}="分包费",{ln}="机械运输费",{ln}="其他直接费",{ln}="票据贴息",{ln}="应收账款")'
        od = f'AND(ISNUMBER({g(J_DATE)}),{g(J_DATE)}<{OPEN_DATE})'
        f[J_CHK] = (f'=IF({empty},"",'
                    f'IF({g(J_ACC)}="","✗ 没选账户",IF({g(J_ATYPE)}="","✗ 账户不在【基础资料】② 里",'
                    f'IF(NOT(ISNUMBER({g(J_DATE)})),"✗ 日期要填成日期（2026/5/1）",IF({od},"✗ 日期早于建账日期（建账前的放【期初余额】）",'
                    f'IF(AND(N({g(J_IN)})<>0,N({g(J_OUT)})<>0),"✗ 收入、支出只填一边",'
                    f'IF({g(J_NET)}=0,"✗ 没有金额",'
                    f'IF({ct}="","✗ 认不出收支类别：在「改类别」选",IF({g(J_ALLOC)}="？","✗ 收支类别不在【基础资料】③ 里",'
                    f'IF(OR({ct}="管理人员工资",{ct}="折旧费"),"✗ 这一类是自动算的，流水里不能选（发工资选「工资」，买设备选「固定资产购置」）",'
                    f'IF(AND({direct},{g(J_PJ)}=""),"✗ 这一笔要选项目（摘要里的项目没在【项目档案】的话先去加上）",'
                    f'IF(AND({g(J_PJ)}<>"",COUNTIF({PJ_NAMES},{g(J_PJ)})=0),"✗ 项目不在【项目档案】里",'
                    f'IF(AND(OR({need}="人",{need}="单位"),{un}=""),"✗ 这一笔要选是谁（改单位/人）",'
                    f'IF(AND({ct}="付应付款",NOT(OR({ut}="材料供应商",{ut}="分包",{ut}="机械运输",{ut}="其他"))),'
                    f'"✗ 付应付款要选材料商、分包或机械单位",'
                    f'IF(AND({ct}="付应付款",{g(J_HANG)}=0),"✗ 这家没有应付：先在【应付登记】登记（现买现付的选材料费等）",'
                    f'IF(AND(OR({ct}="工资",{ct}="工资退回"),{g(J_WKIND)}=0),"✗ 工资要选管理人员或工人（临时工选「临时工工资」）",'
                    f'IF(AND({ct}="账户互转",{g(J_TO)}=""),"✗ 账户互转要在「对方账户」选（借款、报销选那个人的个人户）",'
                    f'IF(AND({ct}="账户互转",{g(J_TOTYPE)}=""),"✗ 对方账户不在【基础资料】② 里",'
                    f'IF(AND({ct}="账户互转",{g(J_TO)}={g(J_ACC)}),"✗ 对方账户不能是自己",'
                    f'IF(IF({ct}="账户互转",COUNTIFS({jr(J_ACC)},{g(J_TO)},{jr(J_TO)},{g(J_ACC)},{jr(J_CT)},"账户互转",{jr(J_NET)},-{g(J_NET)})>0,FALSE),'
                    f'"✗ 对方账户那边也记了这笔互转（只记一边：清空其中一行的内容）",'
                    f'IF({g(J_PKN)}>0,"⚠ 摘要里有两个项目：拆成几行，或在「改项目」选",'
                    f'IF(AND({ct}<>"账户互转",{g(J_TOACC)}<>""),"⚠ 填了对方账户但类别不是账户互转（对方账户不起作用）",'
                    f'IF(AND({un}<>"",{ut}=""),"⚠ 单位/人没在【往来单位】登记",'
                    f'IF(AND(N({g(J_OUT)})<0,{g(J_IO)}="收"),"⚠ 这是收款，以后请记在收入栏",'
                    f'IF(N({g(J_IN)})<0,"⚠ 收入填了负数，已按支出算",'
                    f'IF(N({g(J_OUT)})<0,"↩ 冲回（退款）",'
                    f'IF(AND({g(J_NET)}<0,{ct}<>"固定资产购置",{ct}<>"付应付款",ISNUMBER({g(J_DATE)}),'
                    f'IFERROR(SUMPRODUCT(ISNUMBER(SEARCH({FA_NAMES},{memo}))*({FA_NAMES}<>"")*(({g(J_DN)}-{FA_DATES})<=30)*(({g(J_DN)}-{FA_DATES})>=-30)),0)>0),'
                    f'"⚠ 像是在买【基础资料】⑥ 登记的车/设备：类别改成「固定资产购置」，不然成本算两遍",'
                    f'IF(AND(OR({ct}="材料费",{ct}="分包费",{ct}="机械运输费",{ct}="其他直接费"),{un}="",{g(J_PJ)}<>"",'
                    f'COUNTIFS({apr(AP_PJ)},{g(J_PJ)},{apr(AP_AMT)},-{g(J_NET)})>0),'
                    f'"⚠ 这个项目有同样金额的应付登记：是付那家的，改类别选「付应付款」、改单位选那家（不然成本算两遍）",'
                    f'"√"))))))))))))))))))))))))))))')
        for col, v in f.items():
            ws[f'{col}{r}'] = v
    cols = [J_SEQ, J_DATE, J_ACC, J_MEMO, J_IN, J_OUT, J_BAL, J_PJ, J_CT, J_UN, J_CHK, J_MPJ, J_MCT, J_MUN, J_TOACC, J_NOTE]
    style_rows(ws, J_R0, J_R1, cols, auto=[J_SEQ, J_BAL, J_PJ, J_CT, J_UN, J_CHK],
               fmts={J_DATE: DATE, J_IN: MONEY, J_OUT: MONEY, J_BAL: MONEY},
               aligns={J_MEMO: ALW, J_CHK: AL, J_IN: AR, J_OUT: AR, J_BAL: AR, J_NOTE: AL},
               fills={c: FILL_IN for c in (J_DATE, J_ACC, J_MEMO, J_IN, J_OUT, J_MPJ, J_MCT, J_MUN, J_TOACC, J_NOTE)})
    for r in range(J_R0, J_R1 + 1):
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hidden])
    dv_date(ws, f'{J_DATE}{J_R0}:{J_DATE}{J_R1}')
    dv_list(ws, f'{J_ACC}{J_R0}:{J_ACC}{J_R1}', f'={AC_NAMES}', '哪个账户（【基础资料】② 登记的）')
    dv_list(ws, f'{J_TOACC}{J_R0}:{J_TOACC}{J_R1}', f'={AC_NAMES}', '账户互转：钱转到（或从）哪个账户；借款、报销选那个人的个人户')
    dv_list(ws, f'{J_MPJ}{J_R0}:{J_MPJ}{J_R1}', f'={PJ_NAMES}', '认错了或没认出才选')
    dv_list(ws, f'{J_MCT}{J_R0}:{J_MCT}{J_R1}', f'={CT_NAMES}', '认错了或没认出才选')
    dv_list(ws, f'{J_MUN}{J_R0}:{J_MUN}{J_R1}', f'={UN_NAMES}', '认错了或没认出才选（代别人付的，选被代付的那家）', stop=False)
    rng_ = f'{J_CHK}{J_R0}:{J_CHK}{J_R1}'
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},1)="✗"'], fill=FILL_WARN,
                                                    font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${J_CHK}{J_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'${J_CHK}{J_R0}="√"'],
                                                    font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    for auto_c, man_c in ((J_PJ, J_MPJ), (J_CT, J_MCT), (J_UN, J_MUN)):
        ws.conditional_formatting.add(f'{auto_c}{J_R0}:{auto_c}{J_R1}', FormulaRule(
            formula=[f'${man_c}{J_R0}<>""'], font=Font(name=YH, sz=10, bold=True, color='FF0070C0')))
    ws.conditional_formatting.add(f'{J_BAL}{J_R0}:{J_BAL}{J_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER(${J_BAL}{J_R0}),${J_BAL}{J_R0}<0,${J_ACC}{J_R0}<>"")'], font=Font(name=YH, sz=10, color='FFC00000')))
    ws.auto_filter.ref = f'A{J_HDR}:{J_NOTE}{J_R1}'
    ws.freeze_panes = f'E{J_R0}'
    for i, row in enumerate(ctx['journal']):
        r = J_R0 + i
        for c, v in zip((J_DATE, J_ACC, J_MEMO, J_IN, J_OUT), row[:5]):
            if v is not None:
                ws[f'{c}{r}'] = v
        for c, v in (row[5] if len(row) > 5 else {}).items():
            ws[f'{c}{r}'] = v
    return ws
