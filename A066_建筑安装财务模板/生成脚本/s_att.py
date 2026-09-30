# -*- coding: utf-8 -*-
"""【考勤工资】每人每月一行：项目＋天数成对填（一行 4 对，一个月跑了 5 个以上工地就再加一行）。
   单价自动按【工资标准】查这个月最后一天有效的日薪/月薪；月中涨薪拆两行，涨薪前那行手填旧单价。
   月薪的人：应发＝月薪，天数只用来拆到各项目（算项目人工）和「公司管理」（算管理费）；天数都不填＝全部算公司。
   应发（含补贴、其他加减）按天数比例拆到各项目，这样 各项目人工 ＋ 公司部分 ＝ 应发合计，一分不差。"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *


def build_att(wb, ctx):
    ws = wb.create_sheet(SH_ATT)
    widths(ws, {AT_SEQ: 5, AT_MON: 9, AT_NAME: 9, **{c: 12 for c in AT_PJS}, **{c: 6 for c in AT_DDS},
                AT_URATE: 9, AT_ALLOW: 8, AT_ADJ: 9, AT_TAX: 8, AT_SOC: 8, AT_NOTE: 14,
                AT_DAYS: 7, AT_RATE: 9, AT_WAY: 6, AT_BASE: 11, AT_PAY: 11, AT_NETPAY: 11, AT_CHK: 34})
    title(ws, '考 勤 工 资（每人每月一行 · 项目＋天数成对填 · 单价自动按工资标准查）', AT_CHK, C_HOME,
          '💡 每个月每人一行：月份（这个月任意一天）、姓名，然后「项目1 天数1、项目2 天数2……」在哪个工地干了几天就填几天（半天填 0.5），'
          '在公司/办公室干的项目选「公司管理」。一个月跑了 5 个以上工地，同一个人再加一行接着填。删行请清空内容，别删整行。'
          '单价自动从【工资标准】查（涨工资在那边加一行，以前的月份不变）；月中涨薪：这个月拆两行，涨薪前的天数那一行在「单价手填」填旧单价。'
          '月薪的人（老板、负责人、办公室）：应发＝月薪，在工地干了几天就填那个项目几天、其余填「公司管理」——按天数把工资拆到项目和公司；不填天数＝全部算公司管理费。'
          '零星的临时工、点工当场结清的不用录这里，直接在【资金流水】记「临时工工资」、选项目。'
          '月薪的人一个月拆成两行（月中涨薪、超过 4 个工地）时，月薪按两行的天数比例分，不会算两遍。'
          '发工资不在这里记：在【资金流水】记（类别「工资」），【工资表】自动算已发多少、还欠多少。')
    heads = [(AT_SEQ, '序号'), (AT_MON, '月份'), (AT_NAME, '姓名')]
    for k in range(4):
        heads += [(AT_PJS[k], f'项目{k + 1}'), (AT_DDS[k], f'天数{k + 1}')]
    heads += [(AT_URATE, '单价手填\n（空＝自动）'), (AT_ALLOW, '补贴'), (AT_ADJ, '其他加减\n（±）'), (AT_TAX, '代扣\n个税'),
              (AT_SOC, '代扣\n社保'), (AT_NOTE, '备注'), (AT_DAYS, '合计\n天数'), (AT_RATE, '单价'), (AT_WAY, '计薪'),
              (AT_BASE, '基本工资'), (AT_PAY, '应发合计'), (AT_NETPAY, '实发应付\n（给他的）'), (AT_CHK, '校验')]
    for col, t in heads:
        colr = ('FF305496' if col in (AT_SEQ, AT_MON, AT_NAME) else
                'FF548235' if col in AT_PJS + AT_DDS else
                'FFBF8F00' if col in (AT_URATE, AT_ALLOW, AT_ADJ, AT_TAX, AT_SOC, AT_NOTE) else 'FF7F7F7F')
        put(ws, f'{col}{AT_HDR}', t, F_HDR, fill(colr), align=ACW)
    ws.row_dimensions[AT_HDR].height = 34
    hid = [(AT_YM, '年月'), (AT_EFF, '生效日'), (AT_SDAY, '标准日薪'), (AT_SMON, '标准月薪')] + \
          [(c, f'项目{i + 1}金额') for i, c in enumerate(AT_AMTS)] + [(AT_COAMT, '公司部分'), (AT_MIDRAISE, '月中调价日'), (AT_KIND, '人员类型')]
    for col, t in hid:
        ws[f'{col}{AT_HDR}'] = t
        ws[f'{col}{AT_HDR}'].font = F_HELP
    for r in range(AT_R0, AT_R1 + 1):
        g = lambda c: f'{c}{r}'
        nm, mon = g(AT_NAME), g(AT_MON)
        blank = f'AND({mon}="",{nm}="")'
        ws[g(AT_SEQ)] = f'=IF({blank},"",ROW()-{AT_HDR})'
        ws[g(AT_YM)] = f'=IF(ISNUMBER({mon}),YEAR({mon})*100+MONTH({mon}),"")'
        ws[g(AT_DAYS)] = f'=IF({nm}="","",' + '+'.join(f'N({g(c)})' for c in AT_DDS) + ')'
        ok = f'AND({nm}<>"",ISNUMBER({mon}))'
        ws[g(AT_EFF)] = f'=IF({ok},SUMPRODUCT(MAX(({RT_NAMES}={nm})*({RT_DNS}<=EOMONTH({mon},0))*{RT_DNS})),0)'
        ws[g(AT_SDAY)] = f'=IF(N({g(AT_EFF)})=0,0,SUMIFS({RT_DAYS},{RT_NAMES},{nm},{RT_DATES},{g(AT_EFF)}))'
        ws[g(AT_SMON)] = f'=IF(N({g(AT_EFF)})=0,0,SUMIFS({RT_MONS},{RT_NAMES},{nm},{RT_DATES},{g(AT_EFF)}))'
        ws[g(AT_WAY)] = f'=IF({nm}="","",IF({g(AT_SMON)}>0,"月薪","日薪"))'
        ws[g(AT_RATE)] = f'=IF({nm}="","",IF({g(AT_URATE)}<>"",N({g(AT_URATE)}),IF({g(AT_WAY)}="月薪",{g(AT_SMON)},{g(AT_SDAY)})))'
        pdays = f'SUMIFS({atr(AT_DAYS)},{atr(AT_NAME)},{nm},{atr(AT_YM)},{g(AT_YM)})'
        first = f'COUNTIFS(${AT_NAME}${AT_R0}:{nm},{nm},${AT_YM}${AT_R0}:{g(AT_YM)},{g(AT_YM)})=1'
        ws[g(AT_BASE)] = (f'=IF({nm}="","",ROUND(IF({g(AT_WAY)}="月薪",IF({pdays}=0,IF({first},N({g(AT_RATE)}),0),'
                          f'N({g(AT_RATE)})*N({g(AT_DAYS)})/{pdays}),N({g(AT_DAYS)})*N({g(AT_RATE)})),2))')
        ws[g(AT_PAY)] = f'=IF({nm}="","",ROUND(N({g(AT_BASE)})+N({g(AT_ALLOW)})+N({g(AT_ADJ)}),2))'
        ws[g(AT_NETPAY)] = f'=IF({nm}="","",ROUND(N({g(AT_PAY)})-N({g(AT_TAX)})-N({g(AT_SOC)}),2))'
        for pc, dc, ac in zip(AT_PJS, AT_DDS, AT_AMTS):
            ws[g(ac)] = f'=IF(OR({g(pc)}="",N({g(dc)})=0,N({g(AT_DAYS)})=0),0,ROUND(N({g(AT_PAY)})*{g(dc)}/{g(AT_DAYS)},2))'
        proj_part = '+'.join(f'{g(ac)}*({g(pc)}<>"{CO_PSEUDO}")' for pc, ac in zip(AT_PJS, AT_AMTS))
        ws[g(AT_COAMT)] = f'=IF({nm}="",0,ROUND(N({g(AT_PAY)})-({proj_part}),2))'
        ws[g(AT_MIDRAISE)] = (f'=IF(NOT({ok}),0,SUMPRODUCT(MAX(({RT_NAMES}={nm})*({RT_DNS}>{mon}-DAY({mon})+1)'
                              f'*({RT_DNS}<=EOMONTH({mon},0))*{RT_DNS})))')
        ws[g(AT_KIND)] = f'=IF({nm}="","",IFERROR(INDEX({UN_TYPES_R},MATCH({nm},{UN_NAMES},0))&"",""))'
        badpj = 'OR(' + ','.join(f'AND({g(pc)}<>"",{g(pc)}<>"{CO_PSEUDO}",COUNTIF({PJ_NAMES},{g(pc)})=0)' for pc in AT_PJS) + ')'
        daynopj = 'OR(' + ','.join(f'AND({g(pc)}="",N({g(dc)})<>0)' for pc, dc in zip(AT_PJS, AT_DDS)) + ')'
        pjnoday = 'OR(' + ','.join(f'AND({g(pc)}<>"",N({g(dc)})=0)' for pc, dc in zip(AT_PJS, AT_DDS)) + ')'
        dup = 'OR(' + ','.join(f'AND({g(AT_PJS[a])}<>"",{g(AT_PJS[a])}={g(AT_PJS[b])})' for a in range(4) for b in range(a + 1, 4)) + ')'
        ws[g(AT_CHK)] = (f'=IF({blank},"",IF(NOT(ISNUMBER({mon})),"✗ 月份要填成日期（这个月任意一天）",'
                         f'IF({nm}="","✗ 填姓名",IF({g(AT_KIND)}="","✗ 这个人没在【往来单位】登记",'
                         f'IF({g(AT_KIND)}="临时工","✗ 临时工现结不走考勤：在【资金流水】记「临时工工资」（要代扣个税的，在往来单位改成「工人」）",'
                         f'IF(AND({g(AT_KIND)}<>"工人",{g(AT_KIND)}<>"管理人员"),"✗ 只有工人、管理人员录考勤（往来单位里的类型）",'
                         f'IF({mon}<{OPEN_DATE}-DAY({OPEN_DATE})+1,"✗ 月份早于建账日期",'
                         f'IF(AND({g(AT_URATE)}="",N({g(AT_RATE)})=0),"✗ 【工资标准】里查不到他这个月的单价（工资算成 0）",'
                         f'IF({badpj},"✗ 项目不在【项目档案】里",IF({daynopj},"✗ 有天数没选项目",IF({pjnoday},"✗ 选了项目没填天数",'
                         f'IF(AND({g(AT_WAY)}="日薪",N({g(AT_DAYS)})=0,N({g(AT_PAY)})<>0),"✗ 日薪的人要填天数和项目",'
                         f'IF({dup},"⚠ 同一个项目填了两次",'
                         f'IF(AND(N({g(AT_MIDRAISE)})>0,{g(AT_URATE)}="",COUNTIFS({atr(AT_NAME)},{nm},{atr(AT_YM)},{g(AT_YM)})=1),'
                         f'"⚠ 本月"&DAY({g(AT_MIDRAISE)})&"号调价：拆两行，调价前那行手填旧单价",'
                         f'IF(SUMIFS({atr(AT_DAYS)},{atr(AT_NAME)},{nm},{atr(AT_YM)},{g(AT_YM)})>DAY(EOMONTH({mon},0)),'
                         f'"⚠ 这个人这个月天数比当月天数还多",'
                         f'"√")))))))))))))))')
    cols = [AT_SEQ, AT_MON, AT_NAME] + [c for pair in zip(AT_PJS, AT_DDS) for c in pair] + \
           [AT_URATE, AT_ALLOW, AT_ADJ, AT_TAX, AT_SOC, AT_NOTE, AT_DAYS, AT_RATE, AT_WAY, AT_BASE, AT_PAY, AT_NETPAY, AT_CHK]
    style_rows(ws, AT_R0, AT_R1, cols, auto=[AT_SEQ, AT_DAYS, AT_RATE, AT_WAY, AT_BASE, AT_PAY, AT_NETPAY, AT_CHK],
               fmts={AT_MON: 'yyyy"年"m"月"', AT_RATE: '#,##0.##', AT_URATE: '#,##0.##', AT_BASE: MONEY, AT_PAY: MONEY,
                     AT_NETPAY: MONEY, AT_ALLOW: MONEY, AT_ADJ: MONEY, AT_TAX: MONEY, AT_SOC: MONEY,
                     **{c: '0.##' for c in AT_DDS + [AT_DAYS]}},
               aligns={AT_NOTE: AL, AT_CHK: AL, **{c: AL for c in AT_PJS}},
               fills={c: FILL_IN for c in [AT_MON, AT_NAME] + AT_PJS + AT_DDS + [AT_URATE, AT_ALLOW, AT_ADJ, AT_TAX, AT_SOC, AT_NOTE]})
    for r in range(AT_R0, AT_R1 + 1):
        for c, _ in hid:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hid])
    rng_ = f'{AT_CHK}{AT_R0}:{AT_CHK}{AT_R1}'
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${AT_CHK}{AT_R0},1)="✗"'], fill=FILL_WARN,
                                                    font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${AT_CHK}{AT_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'${AT_CHK}{AT_R0}="√"'],
                                                    font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    dv_date(ws, f'{AT_MON}{AT_R0}:{AT_MON}{AT_R1}')
    dv_list(ws, f'{AT_NAME}{AT_R0}:{AT_NAME}{AT_R1}', f'={UN_NAMES}', '【往来单位】登记过的人')
    for pc in AT_PJS:
        dv_list(ws, f'{pc}{AT_R0}:{pc}{AT_R1}', f'={AX_ALL_R}', '项目，或者「公司管理」（在公司/办公室干的）')
    dv = DataValidation(type='decimal', operator='between', formula1='0', formula2='31', allow_blank=True,
                        showErrorMessage=True, errorTitle='天数', error='填 0～31 的数（半天填 0.5）')
    ws.add_data_validation(dv)
    for dc in AT_DDS:
        dv.add(f'{dc}{AT_R0}:{dc}{AT_R1}')
    ws.freeze_panes = f'D{AT_R0}'
    ws.auto_filter.ref = f'A{AT_HDR}:{AT_CHK}{AT_R1}'
    for i, row in enumerate(ctx['att_rows']):
        r = AT_R0 + i
        ws[f'{AT_MON}{r}'], ws[f'{AT_NAME}{r}'] = row['mon'], row['name']
        for k, (pj, dd) in enumerate(row.get('pairs', [])):
            ws[f'{AT_PJS[k]}{r}'], ws[f'{AT_DDS[k]}{r}'] = pj, dd
        for c, key in ((AT_URATE, 'urate'), (AT_ALLOW, 'allow'), (AT_ADJ, 'adj'), (AT_TAX, 'tax'), (AT_SOC, 'soc'), (AT_NOTE, 'note')):
            if row.get(key) not in (None, ''):
                ws[f'{c}{r}'] = row[key]
    return ws
