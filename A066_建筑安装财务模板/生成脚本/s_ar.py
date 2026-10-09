# -*- coding: utf-8 -*-
"""【应收账款】按项目总表（下面接「按甲方汇总」）、【应收对账】客户对账单（选项目或甲方＋起止：汇总＋三张清单）。
   口径：报表口径.md §3 合同口径——应收总额＝合同（开工日起算）＋补充协议＋签证＋结算调整−扣款＋期初应收；
   已收＝收支登记归类「工程款收款」（含专户、个人户收到的）；已开票＝应收登记类型「开票」。
   只用定义名称；清单用本表右边隐藏列（条件＋计数 / 日期排序键）按第 n 条取数。
   两张表下半部分都是「活动行」：隐藏列先算出每一行该显示什么（项目行 / 甲方行 / 表头 / 合计…），
   所以几块内容一块接一块往下排，中间不留一大片空行；打印区域也跟着内容伸缩。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, Border, Side, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from common import *
from layout import *

NL = 300                                    # 每张清单最多显示多少行
RATES = [('9%', 0.09), ('13%', 0.13), ('6%', 0.06), ('3%', 0.03), ('1%', 0.01)]
_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_WHITE_B = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_GREEN_B = Font(name=YH, sz=10, bold=True, color='FF375623')
F_GRAY = Font(name=YH, sz=9, color='FF808080')
F_BIG = Font(name=YH, sz=15, bold=True, color='FF000000')
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC00000')
EFF_FMT = '"实际 "yyyy-mm-dd'


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def _sw(code_cell, m):
    """按这一行的类型代码选公式：m = {代码: 公式片段}"""
    return '=' + _chain([(f'{code_cell}="{k}"', v) for k, v in m.items()])


def _rate_rng(name, rate):
    return f'{name},">="&{rate - 0.005:.3f},{name},"<"&{rate + 0.005:.3f}'


def _q(text):
    return '"' + text.replace('"', '""') + '"'


def _print_area(wb, ws, ref):
    """打印区域跟着内容伸缩：Print_Area＝固定左上角 : INDEX(最后一列, 最后一行)"""
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=ref)


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def _cf(ws, rng, first_code_cell, rules):
    """rules: [(条件公式(用 {c} 代表本行类型格), fill, font, border)]"""
    for cond, fl, ft, bd in rules:
        kw = {}
        if fl is not None:
            kw['fill'] = fl
        if ft is not None:
            kw['font'] = ft
        if bd is not None:
            kw['border'] = bd
        ws.conditional_formatting.add(rng, FormulaRule(formula=[cond.format(c=first_code_cell)], stopIfTrue=True, **kw))


# ═══════════════════════════════ 应收账款（按项目总表） ═══════════════════════════════
A_GRP, A_HDR, A_TOT, A_R0 = 4, 5, 6, 7
A_NREG = N_PJ + 100                         # 活动区行数：项目最多 200 行，后面接按甲方汇总
A_LAST = 'AE'
# 显示列：(列, 表头, 指标键, 种类, 列宽)
A_DISP = [
    ('A', '序号', None, 'seq', 5), ('B', '项目', None, 'name', 18), ('C', '甲方', None, 'jf', 26),
    ('D', '合同金额\n（开工后才算）', '合同', 'm', 14), ('E', '补充协议', '补充', 'm', 12), ('F', '签证', '签证', 'm', 12),
    ('G', '结算调整\n（多退少补）', '结算', 'm', 12), ('H', '扣款\n（减）', '扣款', 'm', 12), ('I', '期初应收\n（建账前）', '期初', 'm', 12),
    ('J', '应收总额', '应收', 'm', 14), ('K', '已收', '已收', 'm', 14), ('L', '未收', '未收', 'm', 14), ('M', '回款\n比例', None, 'ratio', 8),
    ('N', '已开票', '已开票', 'm', 14), ('O', '未开票\n（应收−已开票）', '未开票', 'm', 14), ('P', '开票未收\n（已开票−已收）', '开票未收', 'm', 14),
    ('Q', '销项税', '销项税', 'm', 12),
    ('R', '9%', 'r9', 'm', 12), ('S', '13%', 'r13', 'm', 12), ('T', '6%', 'r6', 'm', 11), ('U', '3%', 'r3', 'm', 11),
    ('V', '1%', 'r1', 'm', 11), ('W', '其他税率', 'rx', 'm', 11), ('X', '主要税率', None, 'mainrate', 9),
    ('Y', '保证金', '保证金', 'm', 12), ('Z', '质保金\n（空＝比例×应收）', '质保金', 'm', 12), ('AA', '质保到期', '质保到期', 'pdate', 11),
    ('AB', '本期收款', '本期收款', 'm', 13), ('AC', '本期开票', '本期开票', 'm', 13), ('AD', '收款笔数\n（截至日前）', '笔数', 'int', 8),
    ('AE', '最后回款日', '最后回款', 'pdate', 11),
]
A_GROUPS = [('A', 'C', '项目'), ('D', 'J', '应收总额＝合同＋补充协议＋签证＋结算调整−扣款＋期初应收'), ('K', 'M', '收款'),
            ('N', 'Q', '开票'), ('R', 'X', '已开票按税率分（含税金额）'), ('Y', 'AA', '保证金 · 质保金'),
            ('AB', 'AC', '本期（本期起～截至）'), ('AD', 'AE', '收款记录')]
# 隐藏列
AK_I, AK_NM, AK_FL, AK_JF, AK_JFK, AK_HAS = 'AK', 'AL', 'AM', 'AN', 'AO', 'AP'
A_METRICS = ['合同', '补充', '签证', '结算', '扣款', '期初', '应收', '已收', '未收', '已开票', '未开票', '开票未收', '销项税',
             'r9', 'r13', 'r6', 'r3', 'r1', 'rx', '保证金', '质保金', '质保到期', '本期收款', '本期开票', '笔数', '最后回款']
A_HC = {k: CL(CI('AQ') + i) for i, k in enumerate(A_METRICS)}        # AQ..BP
A_PCNT = CL(CI('AQ') + len(A_METRICS))                               # BQ 项目计数
A_CNM, A_CCNT = 'BS', 'BT'                                           # 甲方名、甲方计数
A_SI, A_SKEY = 'BV', 'BW'                                            # _收 第 i 条、最后回款排序键
A_K, A_TYP, A_IDX = 'AG', 'AH', 'AI'                                 # 活动区：第 k 行、类型、第几个
A_NP, A_NC, A_NO, A_LASTK, A_LASTR = '$AG$1', '$AG$2', '$AG$3', '$AG$4', '$AG$5'


def build_ar(ws):
    E, S = '$D$3', '$H$3'
    widths(ws, {c: w for c, _h, _k, _t, w in A_DISP})
    ws.column_dimensions['AF'].width = 10
    tip = ('💡 每个工程项目一行，全自动不用填。应收总额＝合同金额＋补充协议＋签证＋结算调整−扣款（＋建账前的期初应收）：合同在【项目档案】，'
           '其余在【应收登记】一笔一行记；已收＝【收支登记】里归类为「工程款收款」、选了这个项目的钱（农民工专户、九安专户、个人账户收到的都算）；'
           '已开票＝【应收登记】类型「开票」。黄格：截至日期（空＝首页截止日，之后的都不算；开工日在它以后的项目合同额先不算）、'
           '本期起（「本期收款」「本期开票」从这天算到截至日）。每一笔收款、每一张发票看【应收对账】。项目下面接着是「按甲方汇总」。')
    title(ws, '应收账款（按项目）', A_LAST, C_VIEW, tip)
    home_link(ws, 'AF1')
    # 选择格
    selector(ws, 'B3', '截至日期', 'C3', None, fmt=DATE)
    dv_date(ws, 'C3')
    put(ws, 'D3', '=IF(ISNUMBER(C3),INT(C3),P_截止)', F_AUTOB, FILL_AUTO, EFF_FMT, AC)
    put(ws, 'E3', '空＝首页截止日', F_NOTE, align=AL, border=False)
    selector(ws, 'F3', '本期起', 'G3', None, fmt=DATE)
    dv_date(ws, 'G3')
    put(ws, 'H3', '=IF(ISNUMBER(G3),INT(G3),P_年初)', F_AUTOB, FILL_AUTO, EFF_FMT, AC)
    put(ws, 'I3', '空＝首页年初', F_NOTE, align=AL, border=False)
    hr0, hr1 = A_R0, A_R0 + N_PJ - 1
    rng = lambda c: f'${c}${hr0}:${c}${hr1}'
    nosk = (f'ROUND(SUMIFS(收_净额,收_归类,"工程款收款",收_日期,"<="&{E})-SUM({rng(A_HC["已收"])}),2)')
    ws.merge_cells('J3:Q3')
    put(ws, 'J3', f'=IF({nosk}<>0,"⚠ 有 "&TEXT({nosk},"#,##0.00")&" 元工程款没选项目（或项目不是工程），没算进上表——到【收支登记】补选项目","")',
        F_WARN, align=AL, border=False)
    ws.merge_cells('R3:X3')
    put(ws, 'R3', f'=IF({A_NP}+{A_NC}+8>{A_NREG},"⚠ 项目、甲方太多，下面「按甲方汇总」没显示全","")', F_WARN, align=AL, border=False)
    ws.row_dimensions[3].height = 22
    # 表头（两行：分组、列名）
    for c1, c2, t in A_GROUPS:
        if c1 != c2:
            ws.merge_cells(f'{c1}{A_GRP}:{c2}{A_GRP}')
        put(ws, f'{c1}{A_GRP}', t, F_HDR, fill('FF70AD47'), align=ACW)
        for i in range(CI(c1) + 1, CI(c2) + 1):
            put(ws, f'{CL(i)}{A_GRP}', None, F_HDR, fill('FF70AD47'))
    ws.row_dimensions[A_GRP].height = 20
    header(ws, A_HDR, [(c, h) for c, h, *_r in A_DISP], C_VIEW, height=44)

    # ── 隐藏：每个项目（项目档案第 i 个）的各项指标 ──
    for c, t in [(AK_I, 'i'), (AK_NM, '项目'), (AK_FL, '工程'), (AK_JF, '甲方'), (AK_JFK, '甲方在表'), (AK_HAS, '有数')] + \
            [(A_HC[k], k) for k in A_METRICS] + [(A_PCNT, '项目计数'), (A_CNM, '甲方'), (A_CCNT, '甲方计数'), (A_SI, '_收第i条'),
                                                (A_SKEY, '回款键'), (A_K, 'k'), (A_TYP, '类型'), (A_IDX, '第几个')]:
        ws[f'{c}{A_HDR}'] = t
        ws[f'{c}{A_HDR}'].font = F_HELP
    H = A_HC
    for i in range(N_PJ):
        r = hr0 + i
        ws[f'{AK_I}{r}'] = i + 1
        ix = lambda nm: f'INDEX({nm},${AK_I}{r})'
        nm = f'${AK_NM}{r}'
        en = esc(nm)
        off = f'${AK_FL}{r}=0'
        upto = f'{{s}},"<="&{E}'
        f = {}
        f[AK_NM] = f'={ix("项目_名称")}&""'
        f[AK_FL] = f'=IF(AND({AK_NM}{r}<>"",{ix("项目_类型")}="工程"),1,0)'
        f[AK_JF] = f'=IF({AK_FL}{r}=1,{ix("项目_甲方")}&"","")'
        f[AK_JFK] = f'=IF({AK_FL}{r}=0,0,IF({AK_JF}{r}="",0,IF(ISNUMBER(MATCH({AK_JF}{r},甲方_名称,0)),1,0)))'
        f[AK_HAS] = (f'=IF({AK_FL}{r}=0,0,IF(OR({H["应收"]}{r}<>0,{H["已收"]}{r}<>0,{H["已开票"]}{r}<>0,'
                     f'{H["保证金"]}{r}<>0),1,0))')
        f[H['合同']] = f'=IF({off},0,IF({ix("项目_开工有效")}<={E},{ix("项目_合同")},0))'
        for k, t in (('补充', '补充协议'), ('签证', '签证'), ('结算', '结算调整'), ('扣款', '扣款')):
            f[H[k]] = f'=IF({off},0,SUMIFS(应_金额,应_项目,{en},应_类型,"{t}",{upto.format(s="应_日期")}))'
        f[H['期初']] = f'=IF({off},0,SUMIFS(期初_金额,期初_类型,"应收账款",期初_项目,{en}))'
        f[H['应收']] = (f'=ROUND({H["合同"]}{r}+{H["补充"]}{r}+{H["签证"]}{r}+{H["结算"]}{r}-{H["扣款"]}{r}+{H["期初"]}{r},2)')
        f[H['已收']] = f'=IF({off},0,SUMIFS(收_净额,收_归类,"工程款收款",收_项目,{en},{upto.format(s="收_日期")}))'
        f[H['未收']] = f'=ROUND({H["应收"]}{r}-{H["已收"]}{r},2)'
        f[H['已开票']] = f'=IF({off},0,SUMIFS(应_开票额,应_项目,{en},{upto.format(s="应_日期")}))'
        f[H['未开票']] = f'=ROUND({H["应收"]}{r}-{H["已开票"]}{r},2)'
        f[H['开票未收']] = f'=ROUND({H["已开票"]}{r}-{H["已收"]}{r},2)'
        f[H['销项税']] = f'=IF({off},0,SUMIFS(应_销项税,应_项目,{en},{upto.format(s="应_日期")}))'
        for (lbl, rt), k in zip(RATES, ('r9', 'r13', 'r6', 'r3', 'r1')):
            f[H[k]] = f'=IF({off},0,SUMIFS(应_开票额,应_项目,{en},{_rate_rng("应_税率", rt)},{upto.format(s="应_日期")}))'
        f[H['rx']] = f'=ROUND({H["已开票"]}{r}-SUM({H["r9"]}{r}:{H["r1"]}{r}),2)'
        f[H['保证金']] = f'=IF({off},0,N({ix("项目_保证金")}))'
        f[H['质保金']] = (f'=IF({off},0,IF(N({ix("项目_质保金额")})<>0,N({ix("项目_质保金额")}),'
                         f'ROUND(N({ix("项目_质保比例")})*{H["应收"]}{r},2)))')
        f[H['质保到期']] = f'=IF({off},0,N({ix("项目_质保到期")}))'
        f[H['本期收款']] = (f'=IF({off},0,SUMIFS(收_净额,收_归类,"工程款收款",收_项目,{en},收_日期,">="&{S},收_日期,"<="&{E}))')
        f[H['本期开票']] = f'=IF({off},0,SUMIFS(应_开票额,应_项目,{en},应_日期,">="&{S},应_日期,"<="&{E}))'
        f[H['笔数']] = f'=IF({off},0,COUNTIFS(收_归类,"工程款收款",收_项目,{en},{upto.format(s="收_日期")}))'
        # 最后回款日：回款键＝项目序号×100000＋日期；比（序号+1）×100000 小的键里最大的那个，属于本项目就是最后一次回款
        kr = f'${A_SKEY}${A_R0}:${A_SKEY}${A_R0 + N_CASH - 1}'
        f[H['最后回款']] = (f'=IF({off},0,IFERROR(MAX(0,SMALL({kr},COUNTIF({kr},"<"&(${AK_I}{r}+1)*100000))-${AK_I}{r}*100000),0))')
        f[A_CNM] = f'=INDEX(甲方_名称,${AK_I}{r})&""' if i < N_CUS else '=""'
        for c, v in f.items():
            ws[f'{c}{r}'] = v
    counter(ws, A_PCNT, hr0, N_PJ, lambda i: f'${AK_FL}{hr0 + i}=1')
    counter(ws, A_CCNT, hr0, N_PJ,
            lambda i: f'AND(${A_CNM}{hr0 + i}<>"",COUNTIFS({rng(AK_JF)},{esc(f"${A_CNM}{hr0 + i}")},{rng(AK_HAS)},1)>0)')
    # 回款键（_收 每一条）
    for i in range(N_CASH):
        r = A_R0 + i
        n = f'${A_SI}{r}'
        ws[f'{A_SI}{r}'] = i + 1
        ws[f'{A_SKEY}{r}'] = (f'=IF(INDEX(收_归类,{n})<>"工程款收款",0,IF(OR(INDEX(收_项目,{n})="",INDEX(收_日期,{n})>{E},'
                              f'INDEX(收_净额,{n})<=0),0,IFERROR(MATCH(INDEX(收_项目,{n}),{rng(AK_NM)},0)*100000+INDEX(收_日期,{n}),0)))')
    # 标量
    ws[A_NP.replace('$', '')] = f'=${A_PCNT}${hr1}'
    ws[A_NC.replace('$', '')] = f'=${A_CCNT}${hr1}'
    ws[A_NO.replace('$', '')] = f'=COUNTIFS({rng(AK_FL)},1,{rng(AK_JFK)},0,{rng(AK_HAS)},1)'
    ws[A_LASTK.replace('$', '')] = f'={A_NP}+6+{A_NC}+IF({A_NO}>0,1,0)'
    ws[A_LASTR.replace('$', '')] = f'={A_R0 - 1}+MIN({A_LASTK},{A_NREG})'

    # ── 合计（固定在表头下面） ──
    T = A_TOT
    fm = {'m': MONEY, 'ratio': PCT, 'int': '0;-0;"-"', 'pdate': DATE, 'seq': '0'}
    al = {'seq': AC, 'name': AL, 'jf': AL, 'm': AR, 'ratio': AC, 'int': AC, 'pdate': AC, 'mainrate': AC}

    def mainrate(r):
        return (f'IF(N($N{r})=0,"",INDEX({{"9%","13%","6%","3%","1%","其他"}},MATCH(MAX($R{r}:$W{r}),$R{r}:$W{r},0))'
                f'&IF(COUNTIF($R{r}:$W{r},"<>0")>1,"等",""))')

    for c, h, k, kind, _w in A_DISP:
        if kind == 'seq':
            v = None
        elif kind == 'name':
            v = '合计'
        elif kind == 'jf':
            v = f'={A_NP}&" 个工程项目"'
        elif kind in ('m', 'int'):
            v = f'=SUM({rng(H[k])})'
        elif kind == 'ratio':
            v = f'=IF(N($J{T})=0,"",$K{T}/$J{T})'
        elif kind == 'mainrate':
            v = '=' + mainrate(T)
        elif k == '最后回款':
            v = f'=IF(MAX({rng(H[k])})=0,"",MAX({rng(H[k])}))'
        else:
            v = None
        put(ws, f'{c}{T}', v, F_AUTOB, FILL_TOT, fm.get(kind), al.get(kind, AC))
    ws.row_dimensions[T].height = 20

    # ── 活动区：项目行 → 共 N 个 → 按甲方汇总（标题、表头、甲方行、其他、合计、共 N 个） ──
    r0, r1 = A_R0, A_R0 + A_NREG - 1
    JF, FL, JFK = rng(AK_JF), rng(AK_FL), rng(AK_JFK)
    for kk in range(A_NREG):
        r = r0 + kk
        k = f'${A_K}{r}'
        ws[f'{A_K}{r}'] = kk + 1
        np_, nc, no = A_NP, A_NC, A_NO
        ws[f'{A_TYP}{r}'] = (f'=IF({k}<={np_},"P",IF({k}={np_}+1,"N",IF({k}={np_}+3,"S",IF({k}={np_}+4,"H",'
                             f'IF({k}<={np_}+4+{nc},IF({k}>{np_}+4,"C",""),IF(AND({no}>0,{k}={np_}+5+{nc}),"O",'
                             f'IF({k}={np_}+5+{nc}+IF({no}>0,1,0),"T",IF({k}={np_}+6+{nc}+IF({no}>0,1,0),"M",""))))))))')
        ws[f'{A_IDX}{r}'] = (f'=IF(${A_TYP}{r}="P",{kth(k, A_PCNT, hr0, N_PJ)},'
                             f'IF(${A_TYP}{r}="C",{kth(f"({k}-{np_}-4)", A_CCNT, hr0, N_PJ)},0))')
        tc = f'${A_TYP}{r}'
        idx = f'${A_IDX}{r}'
        for c, h, key, kind, _w in A_DISP:
            hshort = h.split('\n')[0]
            m = {}
            if kind == 'seq':
                m = {'P': k, 'C': f'{k}-{np_}-4', 'H': '"序号"'}
            elif kind == 'name':
                m = {'P': f'INDEX({rng(AK_NM)},{idx})', 'C': f'INDEX({rng(A_CNM)},{idx})', 'O': '"（没对上甲方）"',
                     'T': '"甲方合计"', 'H': '"甲方"', 'S': '"▼ 按甲方汇总"', 'N': f'"共 "&{np_}&" 个工程项目"',
                     'M': f'"共 "&{nc}&" 个甲方"'}
            elif kind == 'jf':
                m = {'P': f'INDEX({JF},{idx})', 'C': f'COUNTIFS({JF},{esc(f"$B{r}")},{FL},1)&" 个项目"',
                     'O': f'COUNTIFS({JFK},0,{FL},1)&" 个项目"', 'T': f'{np_}&" 个项目"', 'H': '"项目数"',
                     'S': '"同一甲方的项目加在一起"'}
            elif kind in ('m', 'int'):
                hc = rng(H[key])
                m = {'P': f'INDEX({hc},{idx})', 'C': f'SUMIFS({hc},{JF},{esc(f"$B{r}")},{FL},1)',
                     'O': f'SUMIFS({hc},{JFK},0,{FL},1)', 'T': f'{c}${T}', 'H': _q(hshort)}
            elif kind == 'ratio':
                v = f'IF(N($J{r})=0,"",$K{r}/$J{r})'
                m = {'P': v, 'C': v, 'O': v, 'T': v, 'H': '"回款比例"'}
            elif kind == 'mainrate':
                v = mainrate(r)
                m = {'P': v, 'C': v, 'O': v, 'T': v, 'H': '"主要税率"'}
            elif kind == 'pdate':
                hc = rng(H[key])
                m = {'P': f'IF(INDEX({hc},{idx})=0,"",INDEX({hc},{idx}))', 'H': _q(hshort)}
            ws[f'{c}{r}'] = _sw(tc, m)
            cell = ws[f'{c}{r}']
            cell.font = F_AUTO
            cell.alignment = al.get(kind, AC)
            if kind in fm:
                cell.number_format = fm[kind]
    hide(ws, *[CL(i) for i in range(CI('AG'), CI(A_SKEY) + 1)])
    # 活动区样式（条件格式）
    reg = f'A{r0}:{A_LAST}{r1}'
    c0 = f'${A_TYP}{r0}'
    _cf(ws, reg, c0, [
        ('OR({c}="P",{c}="C",{c}="O")', None, None, CF_BD),
        ('{c}="H"', cf_fill(C_VIEW), F_WHITE_B, CF_BD),
        ('{c}="S"', cf_fill('FFE2EFDA'), F_GREEN_B, None),
        ('{c}="T"', cf_fill('FFFCE4D6'), F_AUTOB, CF_BD),
        ('OR({c}="N",{c}="M")', None, F_GRAY, None),
    ])
    ws.freeze_panes = f'C{A_R0}'
    print_setup(ws, f'{A_GRP}:{A_HDR}', landscape=True)
    q = f"'{ws.title}'"
    _print_area(ws.parent, ws, f'{q}!$A$1:INDEX({q}!${A_LAST}$1:${A_LAST}${r1},{q}!{A_LASTR})')


# ═══════════════════════════════ 应收对账（客户对账单） ═══════════════════════════════
D_R0 = 22                                   # 活动区（三张清单）从第 22 行起
D_NREG = 3 * NL + 16
D_LAST = 'H'
D_K, D_CODE, D_J, D_SRC, D_AMT, D_AMT2 = 'N', 'O', 'P', 'Q', 'R', 'S'
D_PI, D_PNM, D_PFL, D_PJF, D_PST, D_PHT, D_PCNT = 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA'
D_YI, D_YK1, D_YK3 = 'AC', 'AD', 'AE'
D_SI, D_SK2 = 'AG', 'AH'
D_SC = [  # 标量（K 列说明、L 列值）
    'P', 'C', '模式', '起', '止', 'n1a', 'n1b', 'n1', 'n2', 'n3', 'b2', 'b3', 'bF', '末行',
    '合计1', '合计1产值', '合计2', '合计3', '合计3税', '项目序号', 'escP', 'escC', '项目甲方']
D_SR = {k: 2 + i for i, k in enumerate(D_SC)}


def build_ard(ws):
    L = lambda k: f'$L${D_SR[k]}'
    P_, C_, MD, S_, E_ = L('P'), L('C'), L('模式'), L('起'), L('止')
    eP, eC = L('escP'), L('escC')
    widths(ws, {'A': 6, 'B': 13, 'C': 18, 'D': 14, 'E': 32, 'F': 15, 'G': 14, 'H': 17, 'I': 10})
    tip = ('💡 给甲方对账用（可以直接打印）。黄格：选「项目」就按这个项目对账；项目空着、只选「甲方」＝这个甲方的全部项目'
           '（按每一笔登记的客户名称，没填客户的按项目档案的甲方）。起、止空着＝首页的年初～截止日。'
           '上面是汇总：期初未收＋本期应收增加−本期已收＝期末未收；截至止日的应收总额、已收、未收、开票和税率。'
           '下面三张清单：应收事项（合同按开工日算）、每一笔收款（日期、金额、哪个账户）、每一张发票（日期、金额、税率、税额、发票号），每张最多显示 300 笔，'
           '清单合计自动跟汇总核对。数据都来自【项目档案】【应收登记】【收支登记】，在那边改。')
    title(ws, '应收对账（客户对账单）', D_LAST, C_VIEW, tip)
    home_link(ws, 'I1')
    # 选择格
    selector(ws, 'B3', '选项目', 'C3', None, '=项目列表', prompt='选了项目就按项目对账（甲方格不起作用）')
    selector(ws, 'D3', '或选甲方', 'E3', None, '=甲方列表', prompt='项目空着时，按这个甲方的全部项目对账')
    selector(ws, 'F3', '起', 'G3', None, fmt=DATE)
    selector(ws, 'F4', '止', 'G4', None, fmt=DATE)
    dv_date(ws, 'G3:G4')
    put(ws, 'H3', '=IF(ISNUMBER(G3),INT(G3),P_年初)', F_AUTOB, FILL_AUTO, EFF_FMT, AC)
    put(ws, 'H4', '=IF(ISNUMBER(G4),INT(G4),P_截止)', F_AUTOB, FILL_AUTO, EFF_FMT, AC)
    ws.merge_cells('B4:E4')
    put(ws, 'B4', '选了项目就按项目；项目空着只选甲方＝这个甲方的全部项目。起止空着＝首页年初～截止日（右边灰格是实际用的日期）',
        F_NOTE, align=ALW, border=False)
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 28
    ws.row_dimensions[5].height = 8
    # 标量
    pidx = L('项目序号')
    sc = {
        'P': '=TRIM($C$3&"")', 'C': '=TRIM($E$3&"")', '模式': f'=IF({P_}<>"",1,IF({C_}<>"",2,0))',
        '起': '=$H$3', '止': '=$H$4',
        'n1a': f'=${D_PCNT}${D_R0 + N_PJ - 1}', 'n1b': f'=${D_YK1}${D_R0 - 1}', 'n1': f'=MIN({NL},{L("n1a")}+{L("n1b")})',
        'n2': f'=MIN({NL},${D_SK2}${D_R0 - 1})', 'n3': f'=MIN({NL},${D_YK3}${D_R0 - 1})',
        'b2': f'={L("n1")}+4', 'b3': f'={L("b2")}+{L("n2")}+4', 'bF': f'={L("b3")}+{L("n3")}+4',
        '末行': f'=MIN({D_R0 - 1}+{L("bF")}+3,{D_R0 - 1 + D_NREG})',
        '项目序号': f'=IF({P_}="",0,IFERROR(MATCH({P_},项目_名称,0),0))',
        'escP': f'={esc(P_)}', 'escC': f'={esc(C_)}',
        '项目甲方': f'=IF({pidx}=0,"",INDEX(项目_甲方,{pidx})&"")',
    }
    reg = lambda c: f'${c}${D_R0}:${c}${D_R0 + D_NREG - 1}'
    for k, (code, col) in {'合计1': ('D1', D_AMT), '合计1产值': ('D1', D_AMT2), '合计2': ('D2', D_AMT), '合计3': ('D3', D_AMT),
                           '合计3税': ('D3', D_AMT2)}.items():
        sc[k] = f'=SUMIFS({reg(col)},{reg(D_CODE)},"{code}")'
    for k in D_SC:
        ws[f'K{D_SR[k]}'] = k
        ws[f'K{D_SR[k]}'].font = F_HELP
        ws[f'L{D_SR[k]}'] = sc[k]
        ws[f'L{D_SR[k]}'].font = F_HELP

    def conds(*cs):
        return ''.join(',' + c for c in cs)

    def ys(sm, *cs):
        return f'IF({MD}=1,SUMIFS({sm},应_项目,{eP}{conds(*cs)}),IF({MD}=2,SUMIFS({sm},应_客户,{eC}{conds(*cs)}),0))'

    def sk(*cs):
        return (f'IF({MD}=1,SUMIFS(收_净额,收_归类,"工程款收款",收_项目,{eP}{conds(*cs)}),'
                f'IF({MD}=2,SUMIFS(收_净额,收_归类,"工程款收款",收_客户,{eC}{conds(*cs)}),0))')

    def ht(sm, *cs):
        return (f'IF({MD}=1,SUMIFS({sm},项目_名称,{eP},项目_类型,"工程"{conds(*cs)}),'
                f'IF({MD}=2,SUMIFS({sm},项目_甲方,{eC},项目_类型,"工程"{conds(*cs)}),0))')

    qc = (f'IF({MD}=1,SUMIFS(期初_金额,期初_类型,"应收账款",期初_项目,{eP}),'
          f'IF({MD}=2,SUMIFS(期初_金额,期初_类型,"应收账款",期初_对象,{eC}),0))')
    le = lambda nm, d: f'{nm},"<="&{d}'
    ge = lambda nm, d: f'{nm},">="&{d}'
    bef = f'({S_}-1)'
    # ── 抬头 ──
    ws.merge_cells('A6:H6')
    put(ws, 'A6', '=P_公司&"　应收账款对账单"', F_BIG, align=AC, border=False)
    ws.row_dimensions[6].height = 30
    ws.merge_cells('A7:E7')
    jf = L('项目甲方')
    put(ws, 'A7', f'="对账单位："&IF({MD}=1,IF({jf}="","（项目档案没填甲方）",{jf})&"（项目："&{P_}&"）",IF({MD}=2,{C_},"（还没选）"))',
        F_TXTB, align=AL, border=False)
    ws.merge_cells('F7:H7')
    put(ws, 'F7', f'="对账期间："&TEXT({S_},"yyyy-mm-dd")&" 至 "&TEXT({E_},"yyyy-mm-dd")', F_TXTB, align=AR, border=False)
    ws.merge_cells('A8:H8')
    ix = lambda nm: f'INDEX({nm},{pidx})'
    zb = f'IF(N({ix("项目_质保金额")})<>0,N({ix("项目_质保金额")}),ROUND(N({ix("项目_质保比例")})*$B$18,2))'
    info1 = (f'IF(OR({pidx}=0,{ix("项目_类型")}<>"工程"),"✗ 「"&{P_}&"」不是【项目档案】里的工程项目",'
             f'"项目："&{P_}&"　合同金额 "&TEXT(N({ix("项目_合同")}),"#,##0.00")&"　开工 "&IF(N({ix("项目_开工")})>0,'
             f'TEXT({ix("项目_开工")},"yyyy-mm-dd"),"没填（按建账日）")&"　保证金 "&TEXT(N({ix("项目_保证金")}),"#,##0.00")'
             f'&"　质保金 "&TEXT({zb},"#,##0.00")&IF(N({ix("项目_质保到期")})>0,"（"&TEXT({ix("项目_质保到期")},"yyyy-mm-dd")&" 到期）","")'
             f'&IF(AND({C_}<>"",{C_}<>{jf}),"　（已按项目对账，甲方格不起作用）",""))')
    info2 = (f'"甲方："&{C_}&"　名下工程项目 "&COUNTIFS(项目_甲方,{eC},项目_类型,"工程")&" 个，合同金额合计 "'
             f'&TEXT(SUMIFS(项目_合同,项目_甲方,{eC},项目_类型,"工程"),"#,##0.00")&"　保证金合计 "'
             f'&TEXT(SUMIFS(项目_保证金,项目_甲方,{eC},项目_类型,"工程"),"#,##0.00")')
    put(ws, 'A8', f'=IF({MD}=1,{info1},IF({MD}=2,{info2},"← 先在上面黄格选一个项目，或者选一个甲方"))'
                  f'&IF({S_}>{E_},"　✗ 起 晚于 止","")', F_TXT, align=ALW, border=False)
    ws.row_dimensions[8].height = 30
    # ── 一、本期发生 ──
    section(ws, 9, 'A', 'H', f'="一、本期发生（"&TEXT({S_},"yyyy-mm-dd")&" 至 "&TEXT({E_},"yyyy-mm-dd")&"）"', C_VIEW)
    hdr = lambda row, items: header(ws, row, [('A', '')] + items, C_VIEW, height=34)
    val = lambda cell, f, fmt=MONEY, font=F_AUTOB: put(ws, cell, f, font, None, fmt, AR)
    hdr(10, [('B', '期初未收\n（起日前一天）'), ('C', '＋本期应收增加'), ('D', '−本期已收'), ('E', '＝期末未收'), ('F', '本期开票'),
             ('G', '本期销项税'), ('H', '核对')])
    hdr(12, [('B', '应收增加\n其中'), ('C', '本期开工的合同'), ('D', '补充协议'), ('E', '签证'), ('F', '结算调整'), ('G', '扣款（减）'),
             ('H', '确认产值\n（只记录）')])
    hdr(14, [('B', '本期已收\n其中'), ('C', '银行 / 现金'), ('D', '专户\n（农民工等）'), ('E', '个人账户'), ('F', '收款笔数'), ('G', '开票张数'),
             ('H', '')])
    val('B11', f'=ROUND({ht("项目_合同", le("项目_开工有效", bef))}+{qc}+{ys("应_应收额", le("应_日期", bef))}-{sk(le("收_日期", bef))},2)')
    per = lambda nm: (ge(nm, S_), le(nm, E_))
    val('C13', '=' + ht('项目_合同', *per('项目_开工有效')))
    for c, t in (('D', '补充协议'), ('E', '签证'), ('F', '结算调整'), ('G', '扣款')):
        val(f'{c}13', '=' + ys('应_金额', f'应_类型,"{t}"', *per('应_日期')))
    val('H13', '=' + ys('应_产值额', *per('应_日期')))
    val('C11', '=ROUND(C13+D13+E13+F13-G13,2)')
    val('D11', '=' + sk(*per('收_日期')))
    val('E11', '=ROUND(B11+C11-D11,2)')
    val('F11', '=' + ys('应_开票额', *per('应_日期')))
    val('G11', '=' + ys('应_销项税', *per('应_日期')))
    put(ws, 'H11', f'=IF({MD}=0,"",IF(ROUND(E11-D18,2)=0,"✓ 期末＝累计未收","✗ 差 "&TEXT(E11-D18,"#,##0.00")))',
        F_AUTOB, None, None, AC)
    val('D15', '=' + sk(*per('收_日期'), '收_账户类型,"专户"'))
    val('E15', '=' + sk(*per('收_日期'), '收_账户类型,"个人户"'))
    val('C15', '=ROUND(D11-D15-E15,2)')
    val('F15', f'=${D_SK2}${D_R0 - 1}', '0;-0;"-"')
    val('G15', f'=${D_YK3}${D_R0 - 1}', '0;-0;"-"')
    for c in 'ABCDEFGH':
        for rr in (11, 13, 15):
            ws[f'{c}{rr}'].border = BD
            if ws[f'{c}{rr}'].value is None:
                ws[f'{c}{rr}'].fill = FILL_AUTO
    # ── 二、截至止日累计 ──
    section(ws, 16, 'A', 'H', f'="二、截至 "&TEXT({E_},"yyyy-mm-dd")&" 累计（从开工算起）"', C_VIEW)
    hdr(17, [('B', '应收总额'), ('C', '已收总额'), ('D', '未收总额'), ('E', '已开票'), ('F', '未开票'), ('G', '销项税额'), ('H', '回款比例')])
    hdr(19, [('B', '已开票\n按税率')] + [(c, lbl) for c, (lbl, _r) in zip('CDEFG', RATES)] + [('H', '其他税率')])
    upto = lambda nm: le(nm, E_)
    val('B18', f'=ROUND({ht("项目_合同", upto("项目_开工有效"))}+{qc}+{ys("应_应收额", upto("应_日期"))},2)')
    val('C18', '=' + sk(upto('收_日期')))
    val('D18', '=ROUND(B18-C18,2)')
    val('E18', '=' + ys('应_开票额', upto('应_日期')))
    val('F18', '=ROUND(B18-E18,2)')
    val('G18', '=' + ys('应_销项税', upto('应_日期')))
    put(ws, 'H18', '=IF(N(B18)=0,"",C18/B18)', F_AUTOB, None, PCT, AC)
    for c, (lbl, rt) in zip('CDEFG', RATES):
        val(f'{c}20', '=' + ys('应_开票额', _rate_rng('应_税率', rt), upto('应_日期')))
    val('H20', '=ROUND(E18-SUM(C20:G20),2)')
    for c in 'ABCDEFGH':
        for rr in (18, 20):
            ws[f'{c}{rr}'].border = BD
            if ws[f'{c}{rr}'].value is None:
                ws[f'{c}{rr}'].fill = FILL_AUTO
    ws.row_dimensions[21].height = 10

    # ── 隐藏：清单的条件和排序键 ──
    for c, t in ((D_K, 'k'), (D_CODE, '类型'), (D_J, '第几笔'), (D_SRC, '源第几条'), (D_AMT, '金额'), (D_AMT2, '金额2'), (D_PI, 'i'),
                 (D_PNM, '项目'), (D_PFL, '工程'), (D_PJF, '甲方'), (D_PST, '开工'), (D_PHT, '合同'), (D_PCNT, '合同计数'),
                 (D_YI, '_应第i条'), (D_YK1, '应收事项键'), (D_YK3, '发票键'), (D_SI, '_收第i条'), (D_SK2, '收款键')):
        ws[f'{c}{D_R0 - 2}'] = t
        ws[f'{c}{D_R0 - 2}'].font = F_HELP
    pr = lambda c: f'${c}${D_R0}:${c}${D_R0 + N_PJ - 1}'
    for i in range(N_PJ):
        r = D_R0 + i
        ii = f'${D_PI}{r}'
        ws[f'{D_PI}{r}'] = i + 1
        ws[f'{D_PNM}{r}'] = f'=INDEX(项目_名称,{ii})&""'
        ws[f'{D_PFL}{r}'] = f'=IF(AND({D_PNM}{r}<>"",INDEX(项目_类型,{ii})="工程"),1,0)'
        ws[f'{D_PJF}{r}'] = f'=INDEX(项目_甲方,{ii})&""'
        ws[f'{D_PST}{r}'] = f'=N(INDEX(项目_开工有效,{ii}))'
        ws[f'{D_PHT}{r}'] = f'=N(INDEX(项目_合同,{ii}))'
    counter(ws, D_PCNT, D_R0, N_PJ,
            lambda i: (f'AND(${D_PFL}{D_R0 + i}=1,${D_PHT}{D_R0 + i}<>0,${D_PST}{D_R0 + i}>={S_},${D_PST}{D_R0 + i}<={E_},'
                       f'IF({MD}=1,${D_PNM}{D_R0 + i}={P_},IF({MD}=2,${D_PJF}{D_R0 + i}={C_},FALSE)))'))
    # _应：应收事项（类型≠开票）、发票（类型＝开票）；_收：工程款收款 —— 键＝日期×10000＋第几条（按日期排）
    for i in range(N_AR):
        r = D_R0 + i
        n = f'${D_YI}{r}'
        ws[f'{D_YI}{r}'] = i + 1
        dt = f'INDEX(应_日期,{n})'
        who = f'IF({MD}=1,INDEX(应_项目,{n})={P_},IF({MD}=2,INDEX(应_客户,{n})={C_},FALSE))'
        base = f'INDEX(应_有效,{n})=1,{dt}>={S_},{dt}<={E_},{who}'
        ws[f'{D_YK1}{r}'] = f'=IF(AND({base},INDEX(应_类型,{n})<>"开票"),{dt}*10000+{n},"")'
        ws[f'{D_YK3}{r}'] = f'=IF(AND({base},INDEX(应_类型,{n})="开票"),{dt}*10000+{n},"")'
    for c in (D_YK1, D_YK3):
        ws[f'{c}{D_R0 - 1}'] = f'=COUNT({c}{D_R0}:{c}{D_R0 + N_AR - 1})'
    for i in range(N_CASH):
        r = D_R0 + i
        n = f'${D_SI}{r}'
        ws[f'{D_SI}{r}'] = i + 1
        dt = f'INDEX(收_日期,{n})'
        who = f'IF({MD}=1,INDEX(收_项目,{n})={P_},IF({MD}=2,INDEX(收_客户,{n})={C_},FALSE))'
        ws[f'{D_SK2}{r}'] = (f'=IF(INDEX(收_归类,{n})<>"工程款收款","",IF(AND({dt}>={S_},{dt}<={E_},{who}),{dt}*10000+{n},""))')
    ws[f'{D_SK2}{D_R0 - 1}'] = f'=COUNT({D_SK2}{D_R0}:{D_SK2}{D_R0 + N_CASH - 1})'
    for c in (D_YK1, D_YK3, D_SK2):
        ws[f'{c}{D_R0 - 1}'].font = F_HELP

    # ── 活动区：① 应收事项 ② 每笔收款 ③ 每张发票（标题、表头、明细、合计）＋签章 ──
    n1a, n1, n2, n3, b2, b3, bF = (L(x) for x in ('n1a', 'n1', 'n2', 'n3', 'b2', 'b3', 'bF'))
    kc = lambda r: f'${D_K}{r}'

    def lst(no, x, n):
        return f'IF({x}=1,"S{no}",IF({x}=2,"H{no}",IF({x}<={n}+2,"D{no}",IF({x}={n}+3,"T{no}",""))))'

    ksort = lambda kx, col, n: ksorted(kx, col, D_R0, n, f'${col}${D_R0 - 1}')
    rate_txt = lambda x: (f'IF({x}=0,"",IF(ROUND({x}*100,1)=ROUND({x}*100,0),TEXT({x},"0%"),TEXT({x},"0.0%")))')
    chk = lambda tot, ref, what: (f'IF(ROUND({tot}-{ref},2)=0,"✓ 跟上面「{what}」一致","✗ 跟上面「{what}」差 "&TEXT({tot}-{ref},"#,##0.00"))')
    more = lambda cnt, txt: f'IF({cnt}>{NL},"超过 {NL} 笔，只显示前 {NL} 笔，请缩短日期",{_q(txt)})'
    inv_no = inref(SH_AR_IN, AR_COLS['发票号'])
    for kk in range(D_NREG):
        r = D_R0 + kk
        k = kc(r)
        ws[f'{D_K}{r}'] = kk + 1
        ws[f'{D_CODE}{r}'] = (f'=IF({k}>{bF},IF({k}-{bF}=1,"F1",IF({k}-{bF}=2,"F2",IF({k}-{bF}=3,"F3",""))),'
                              f'IF({k}>{b3},{lst(3, f"({k}-{b3})", n3)},IF({k}>{b2},{lst(2, f"({k}-{b2})", n2)},{lst(1, k, n1)})))')
        cd = f'${D_CODE}{r}'
        j = f'${D_J}{r}'
        q = f'${D_SRC}{r}'
        ws[f'{D_J}{r}'] = f'=IF(LEFT({cd},1)="D",{k}-IF({cd}="D1",0,IF({cd}="D2",{b2},{b3}))-2,0)'
        ws[f'{D_SRC}{r}'] = (f'=IF({j}=0,0,IF({cd}="D1",IF({j}<={n1a},{kth(j, D_PCNT, D_R0, N_PJ)},{ksort(f"{j}-{n1a}", D_YK1, N_AR)}),'
                             f'IF({cd}="D2",{ksort(j, D_SK2, N_CASH)},{ksort(j, D_YK3, N_AR)})))')
        isc = f'{j}<={n1a}'
        ws[f'{D_AMT}{r}'] = (f'=IF({q}=0,0,IF({cd}="D1",IF({isc},INDEX({pr(D_PHT)},{q}),INDEX(应_应收额,{q})),'
                             f'IF({cd}="D2",INDEX(收_净额,{q}),INDEX(应_开票额,{q}))))')
        ws[f'{D_AMT2}{r}'] = (f'=IF({q}=0,0,IF({cd}="D1",IF({isc},0,INDEX(应_产值额,{q})),IF({cd}="D3",INDEX(应_销项税,{q}),0)))')
        typ1 = f'INDEX(应_类型,{q})'
        disp = {
            'A': {'D1': j, 'D2': j, 'D3': j, 'H1': '"序号"', 'H2': '"序号"', 'H3': '"序号"'},
            'B': {'D1': f'IF({isc},INDEX({pr(D_PST)},{q}),INDEX(应_日期,{q}))', 'D2': f'INDEX(收_日期,{q})', 'D3': f'INDEX(应_日期,{q})',
                  'H1': '"日期"', 'H2': '"收款日期"', 'H3': '"开票日期"', 'S1': '"① 应收事项"', 'S2': '"② 每笔收款"',
                  'S3': '"③ 每张发票"', 'T1': '"合计"', 'T2': '"合计"', 'T3': '"合计"', 'F1': '"我方（盖章）："',
                  'F2': '"经办人："', 'F3': '"日期："'},
            'C': {'D1': f'IF({isc},INDEX({pr(D_PNM)},{q}),INDEX(应_项目,{q}))', 'D2': f'INDEX(收_项目,{q})', 'D3': f'INDEX(应_项目,{q})',
                  'H1': '"项目"', 'H2': '"项目"', 'H3': '"项目"', 'S1': f'"共 "&({n1a}+{L("n1b")})&" 笔"',
                  'S2': f'"共 "&${D_SK2}${D_R0 - 1}&" 笔"', 'S3': f'"共 "&${D_YK3}${D_R0 - 1}&" 张"'},
            'D': {'D1': f'IF({isc},"合同",{typ1})', 'D2': f'INDEX(收_账户,{q})', 'D3': f'INDEX({inv_no},{q}+1)&""',
                  'H1': '"类型"', 'H2': '"收款账户"', 'H3': '"发票号码"'},
            'E': {'D1': (f'IF({isc},"合同金额（按开工日算）",IF({typ1}="确认产值","甲方确认的产值，只记录、不加应收",'
                         f'IF({typ1}="扣款","扣款，减少应收","")))'),
                  'D2': f'INDEX(收_摘要,{q})', 'D3': rate_txt(f'INDEX(应_税率,{q})'),
                  'H1': '"说明"', 'H2': '"摘要"', 'H3': '"税率"',
                  'S1': more(f'({n1a}+{L("n1b")})', '合同、补充协议、签证、结算、扣款'),
                  'S2': more(f'${D_SK2}${D_R0 - 1}', '收支登记里归类「工程款收款」的'),
                  'S3': more(f'${D_YK3}${D_R0 - 1}', '应收登记里类型「开票」的'),
                  'T1': chk(L('合计1'), '$C$11', '本期应收增加'), 'T2': chk(L('合计2'), '$D$11', '本期已收'),
                  'T3': (f'IF(AND(ROUND({L("合计3")}-$F$11,2)=0,ROUND({L("合计3税")}-$G$11,2)=0),"✓ 开票、税额都跟上面一致",'
                         f'"✗ 跟上面差 "&TEXT({L("合计3")}-$F$11,"#,##0.00")&" / 税 "&TEXT({L("合计3税")}-$G$11,"#,##0.00"))'),
                  'F1': '"对方（盖章）："', 'F2': '"经办人："', 'F3': '"日期："'},
            'F': {'D1': f'${D_AMT}{r}', 'D2': f'${D_AMT}{r}', 'D3': f'${D_AMT}{r}', 'T1': L('合计1'), 'T2': L('合计2'), 'T3': L('合计3'),
                  'H1': '"应收金额"', 'H2': '"收款金额"', 'H3': '"开票金额（含税）"'},
            'G': {'D1': f'IF(OR({isc},${D_AMT2}{r}=0),"",${D_AMT2}{r})', 'D2': f'INDEX(收_账户类型,{q})', 'D3': f'${D_AMT2}{r}',
                  'T1': f'IF({L("合计1产值")}=0,"",{L("合计1产值")})', 'T3': L('合计3税'),
                  'H1': '"确认产值"', 'H2': '"账户类型"', 'H3': '"销项税额"'},
            'H': {'D1': f'IF({isc},"项目档案","应收登记第"&INDEX(应_录入行,{q})&"行")', 'D2': f'"收支登记第"&INDEX(收_录入行,{q})&"行"',
                  'D3': f'"应收登记第"&INDEX(应_录入行,{q})&"行"', 'H1': '"在哪登记"', 'H2': '"在哪登记"', 'H3': '"在哪登记"'},
        }
        fmts = {'A': '0', 'B': DATE, 'F': MONEY, 'G': MONEY}
        aligns = {'A': AC, 'B': AC, 'C': AL, 'D': AL, 'E': AL, 'F': AR, 'G': AR, 'H': AL}
        for c, m in disp.items():
            ws[f'{c}{r}'] = _sw(cd, m)
            cell = ws[f'{c}{r}']
            cell.font = F_AUTO
            cell.alignment = aligns[c]
            if c in fmts:
                cell.number_format = fmts[c]
    rg = f'A{D_R0}:{D_LAST}{D_R0 + D_NREG - 1}'
    c0 = f'${D_CODE}{D_R0}'
    _cf(ws, rg, c0, [
        ('LEFT({c},1)="D"', None, None, CF_BD),
        ('LEFT({c},1)="H"', cf_fill(C_VIEW), F_WHITE_B, CF_BD),
        ('LEFT({c},1)="S"', cf_fill('FFE2EFDA'), F_GREEN_B, None),
        ('LEFT({c},1)="T"', cf_fill('FFFCE4D6'), F_AUTOB, CF_BD),
        ('LEFT({c},1)="F"', None, F_TXTB, None),
    ])
    hide(ws, *[CL(i) for i in range(CI('J'), CI(D_SK2) + 1)])
    print_setup(ws, '6:7', landscape=False)
    q = f"'{ws.title}'"
    _print_area(ws.parent, ws, f'{q}!$A$6:INDEX({q}!$H$1:$H${D_R0 + D_NREG - 1},{q}!{L("末行")})')


def build(wb, ctx=None):
    build_ar(wb[SH_AR])
    build_ard(wb[SH_ARD])
