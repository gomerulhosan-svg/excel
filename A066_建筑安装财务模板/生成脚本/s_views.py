# -*- coding: utf-8 -*-
"""查看用的表（全是公式，不用填）：【账户明细】【资金报表】（周报 / 月报 / 自定义）【应收账款总表】【应付账款总表】【个人往来】"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *
import s_agg

ACC_CAP = 1000          # 账户明细最多列多少笔
STMT_CAP = 150          # 对账单每边最多多少笔
APS_ROWS = 120          # 应付总表最多多少家
APS_PJ = 12             # 应付总表最多几个项目列
PER_CAP = 400           # 个人往来明细最多多少笔

GREY = Font(name=YH, sz=10, color='FFBFBFBF')


def jc(col, r):
    return f'{SH_CASH}!${col}${r}'


def dates(s, e, col=J_DATE, rngf=jr):
    return f'{rngf(col)},">="&{s},{rngf(col)},"<="&{e}'


def _sheet(wb, name, color, text, tip, last, wmap, home_col=None):
    ws = wb.create_sheet(name)
    widths(ws, wmap)
    title(ws, text, last, color, tip)
    home_link(ws, f'{home_col or CL(CI(last) + 1)}1')
    ws.column_dimensions[home_col or CL(CI(last) + 1)].width = max(ws.column_dimensions[home_col or CL(CI(last) + 1)].width or 0, 10)
    return ws


def _lbl(ws, coord, text, fill_='FFD9E1F2'):
    put(ws, coord, text, F_KPI_L, fill(fill_), align=AC)


def _val(ws, coord, f, fmt=MONEY, font=F_AUTOB, fill_=FILL_AUTO):
    put(ws, coord, f, font, fill_, fmt, AR if fmt in (MONEY, MONEY0) else AC)


# ═══════════════════════════ 账户明细 ═══════════════════════════
def build_acc(wb, ctx):
    ws = _sheet(wb, SH_ACC, C_CASH, '账 户 明 细（选一个账户 · 看一段时间的每一笔和结存）',
                '💡 黄格子选账户（银行、现金、老板个人户、承兑汇票都行）和起止日期（空着＝本年 1 月 1 日到截止日）。'
                '下面自动列出这个账户这段时间的每一笔：自己记的，还有在别的账户记的「账户互转」（只记一边的那种，这里自动反过来显示）。'
                '期末余额跟银行 App 对不上，就在【基础资料】② 填实际余额和日期，看差额从哪天开始的。',
                'L', {'A': 6, 'B': 11, 'C': 40, 'D': 13, 'E': 13, 'F': 14, 'G': 13, 'H': 13, 'I': 16, 'J': 13, 'K': 8, 'L': 20})
    first_acc = ctx['accounts'][0][0]
    _lbl(ws, 'B3', '账户')
    put(ws, 'C3', first_acc, F_SEL, FILL_SEL, align=AC)
    dv_list(ws, 'C3', f'={AC_NAMES}', '选账户')
    _lbl(ws, 'D3', '起')
    put(ws, 'E3', None, F_SEL, FILL_SEL, DATE, AC)
    _lbl(ws, 'F3', '止')
    put(ws, 'G3', None, F_SEL, FILL_SEL, DATE, AC)
    put(ws, 'H3', '（起止空着＝本年初到截止日）', F_NOTE, border=False)
    dv_date(ws, 'E3')
    dv_date(ws, 'G3')
    S, E, A = '$C$4', '$E$4', '$C$3'
    _lbl(ws, 'B4', '期间')
    _val(ws, 'C4', f'=MAX(IF(ISNUMBER(E3),E3,DATE({AX_Y},1,1)),{OPEN_DATE})', DATE, F_AUTO)
    put(ws, 'D4', '～', F_TXT, align=AC)
    _val(ws, 'E4', f'=IF(ISNUMBER(G3),G3,{AX_E})', DATE, F_AUTO)
    ob = (f'IFERROR(INDEX({AC_OPENS},MATCH({A},{AC_NAMES},0)),0)'
          f'+SUMIFS({jr(J_NET)},{jr(J_ACC)},{A},{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<"&{S})'
          f'-SUMIFS({jr(J_NET)},{jr(J_TO)},{A},{jr(J_LINE)},"账户互转",{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<"&{S})')
    R0 = 8
    R1 = R0 + ACC_CAP - 1
    for c, t, f in (('G4', '期初余额', f'=ROUND({ob},2)'), ('I4', '本期收入', f'=SUM(D{R0}:D{R1})'), ('K4', '本期支出', f'=SUM(E{R0}:E{R1})')):
        _lbl(ws, c, t)
        _val(ws, CL(CI(c[0]) + 1) + '4', f)
    _lbl(ws, 'G5', '期末余额')
    _val(ws, 'H5', '=ROUND(H4+J4-L4,2)', font=F_KPI_V)
    real = lambda col: f'IFERROR(INDEX({rng(SH_BASE, col, AC_R0, AC_R1)},MATCH({A},{AC_NAMES},0)),"")'
    _lbl(ws, 'I5', '银行App余额')
    _val(ws, 'J5', f'=IF({real(AC_REAL)}="","",{real(AC_REAL)})')
    _lbl(ws, 'K5', '核对差额')
    _val(ws, 'L5', f'=IF({real(AC_DIFF)}="","",{real(AC_DIFF)})')
    # 隐藏排序键（跟资金流水一行对一行）：日期×10000＋行号，按日期先后列出
    HC = 'Z'
    n = J_R1 - J_R0 + 1
    NC = skey(ws, HC, J_R0, n, lambda i: (
        f'AND(ISNUMBER({jc(J_DATE, J_R0 + i)}),{jc(J_DATE, J_R0 + i)}>={S},{jc(J_DATE, J_R0 + i)}<={E},'
        f'OR({jc(J_ACC, J_R0 + i)}={A},AND({jc(J_TO, J_R0 + i)}={A},{jc(J_LINE, J_R0 + i)}="账户互转")))'),
        lambda i: f'INT({jc(J_DATE, J_R0 + i)})')
    put(ws, 'I3', f'=IF({NC}>{ACC_CAP},"⚠ 这段时间有"&{NC}&"笔，只列前{ACC_CAP}笔，把日期缩短一点","")',
        F_RED, border=False)
    ws.merge_cells('I3:L3')
    header(ws, R0 - 1, [('A', '序号'), ('B', '日期'), ('C', '摘要'), ('D', '收入'), ('E', '支出'), ('F', '结存'), ('G', '项目'),
                        ('H', '收支类别'), ('I', '单位 / 人'), ('J', '对方账户'), ('K', '流水\n序号'), ('L', '备注')], C_CASH)
    IDX, MIR = 'N', 'O'
    for k in range(ACC_CAP):
        r = R0 + k
        ws[f'{IDX}{r}'] = f'={ksorted(k + 1, HC, J_R0, n, NC)}'
        ix = f'${IDX}{r}'
        g = lambda col: f'INDEX({jr(col)},{ix})'
        ws[f'{MIR}{r}'] = f'=IF({ix}=0,0,IF({g(J_ACC)}={A},0,1))'
        mi = f'${MIR}{r}'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",{g(J_DATE)})'
        ws[f'C{r}'] = f'=IF({ix}=0,"",{g(J_MEMO)}&IF({mi}=1,"（在【"&{g(J_ACC)}&"】记的互转）",""))'
        ws[f'D{r}'] = f'=IF({ix}=0,"",IF({mi}=1,MAX(-{g(J_NET)},0),N({g(J_IN)})))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",IF({mi}=1,MAX({g(J_NET)},0),N({g(J_OUT)})))'
        prev = '$H$4' if k == 0 else f'F{r - 1}'
        ws[f'F{r}'] = f'=IF({ix}=0,"",ROUND(N({prev})+N(D{r})-N(E{r}),2))'
        ws[f'G{r}'] = f'=IF({ix}=0,"",{g(J_PJ)}&"")'
        ws[f'H{r}'] = f'=IF({ix}=0,"",{g(J_CT)}&"")'
        ws[f'I{r}'] = f'=IF({ix}=0,"",{g(J_UN)}&"")'
        ws[f'J{r}'] = f'=IF({ix}=0,"",IF({mi}=1,{g(J_ACC)},{g(J_TO)})&"")'
        ws[f'K{r}'] = f'=IF({ix}=0,"",{ix})'
        ws[f'L{r}'] = f'=IF({ix}=0,"",{g(J_NOTE)}&"")'
        ws[f'{IDX}{r}'].font = ws[f'{MIR}{r}'].font = F_HELP
    style_rows(ws, R0, R1, list('ABCDEFGHIJKL'), auto=list('ABCDEFGHIJKL'),
               fmts={'B': DATE, 'D': MONEY, 'E': MONEY, 'F': MONEY}, aligns={'C': AL, 'D': AR, 'E': AR, 'F': AR, 'I': AL, 'L': AL})
    for r in range(R0, R1 + 1):
        ws[f'A{r}'].fill = ws[f'C{r}'].fill = FILL_NONE
        for c in 'BDEGHIJKL':
            ws[f'{c}{r}'].fill = FILL_NONE
    ws.conditional_formatting.add(f'A{R0}:L{R1}', FormulaRule(formula=[f'$O{R0}=1'], fill=fill('FFEAF1FB')))
    hide(ws, HC, IDX, MIR)
    ws.freeze_panes = f'C{R0}'
    return ws


# ═══════════════════════════ 资金报表（周报 / 月报 / 自定义） ═══════════════════════════
def build_fund(wb, ctx):
    ws = _sheet(wb, '资金报表', C_CASH, '资 金 周 报 / 月 报（各账户余额 · 分段收支 · 回款 · 付款 · 收支分类）',
                '💡 黄格子选「周报」「月报」或「自定义」，再选哪一天（空着＝首页选的截止月份，没选就是最后一笔流水那天）：周报＝那天所在的周一到周日，月报＝那天所在的月；'
                '自定义＝自己填起止日期。① 各账户期初、收入、支出、期末（老板个人户是负数＝公司欠他的）；② 月报按周分、周报按天分；'
                '③ 收了哪些项目的工程款；④ 付给谁最多；⑤ 这段时间的钱按收支类别分。①里含账户之间的互转，②～⑤不含（自己账户之间倒钱不算收支）。',
                'K', {'A': 7, 'B': 18, 'C': 13, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 11, 'J': 12, 'K': 20})
    _lbl(ws, 'B3', '报表')
    put(ws, 'C3', '月报', F_SEL, FILL_SEL, align=AC)
    dv_list(ws, 'C3', '"周报,月报,自定义"')
    _lbl(ws, 'D3', '哪一天')
    put(ws, 'E3', None, F_SEL, FILL_SEL, DATE, AC)
    dv_date(ws, 'E3')
    _lbl(ws, 'F3', '自定义 起')
    put(ws, 'G3', None, F_SEL, FILL_SEL, DATE, AC)
    _lbl(ws, 'H3', '止')
    put(ws, 'I3', None, F_SEL, FILL_SEL, DATE, AC)
    dv_date(ws, 'G3')
    dv_date(ws, 'I3')
    d = f'IF(ISNUMBER($E$3),$E$3,IF(ISNUMBER({SEL_E}),{SEL_E},IF(YEAR({AX_LAST})={AX_Y},{AX_LAST},{AX_E})))'
    _lbl(ws, 'B4', '期间')
    _val(ws, 'C4', f'=IF($C$3="周报",{d}-WEEKDAY({d},2)+1,IF($C$3="自定义",IF(ISNUMBER($G$3),$G$3,DATE({AX_Y},1,1)),DATE(YEAR({d}),MONTH({d}),1)))',
         DATE, F_AUTOB)
    put(ws, 'D4', '～', F_TXT, align=AC)
    _val(ws, 'E4', f'=IF($C$3="周报",$C$4+6,IF($C$3="自定义",IF(ISNUMBER($I$3),$I$3,{AX_E}),EOMONTH($C$4,0)))', DATE, F_AUTOB)
    put(ws, 'F4', f'=IF($C$3="周报",TEXT($C$4,"yyyy年m月d日")&"那周",IF($C$3="月报",TEXT($C$4,"yyyy年m月"),"自定义期间"))&" 资金"&$C$3',
        F_KPI_L, align=AL, border=False)
    ws.merge_cells('F4:I4')
    S, E = '$C$4', '$E$4'
    DR = dates(S, E)
    # ① 各账户
    section(ws, 6, 'A', 'K', '① 各账户余额（本期收入、支出含账户之间的互转）', C_CASH)
    header(ws, 7, [('A', '序号'), ('B', '账户'), ('C', '类型'), ('D', '期初余额'), ('E', '本期收入'), ('F', '本期支出'),
                   ('G', '期末余额'), ('H', '银行App余额'), ('I', '核对日期'), ('J', '核对差额'), ('K', '说明')], C_CASH)
    A0 = 8
    na = AC_R1 - AC_R0 + 1
    for i in range(na):
        r, a = A0 + i, AC_R0 + i
        nm = f'$B{r}'
        ws[f'A{r}'] = f'=IF({nm}="","",{i + 1})'
        ws[f'B{r}'] = f'={SH_BASE}!${AC_NAME}${a}&""'
        ws[f'C{r}'] = f'={SH_BASE}!${AC_TYPE}${a}&""'
        before = (f'N({SH_BASE}!${AC_OPEN}${a})+SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<"&{S})'
                  f'-SUMIFS({jr(J_NET)},{jr(J_TO)},{nm},{jr(J_LINE)},"账户互转",{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<"&{S})')
        ws[f'D{r}'] = f'=IF({nm}="","",ROUND({before},2))'
        ws[f'E{r}'] = (f'=IF({nm}="","",SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_NET)},">0",{DR})'
                       f'-SUMIFS({jr(J_NET)},{jr(J_TO)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},"<0",{DR}))')
        ws[f'F{r}'] = (f'=IF({nm}="","",-SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_NET)},"<0",{DR})'
                       f'+SUMIFS({jr(J_NET)},{jr(J_TO)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},">0",{DR}))')
        ws[f'G{r}'] = f'=IF({nm}="","",ROUND(D{r}+E{r}-F{r},2))'
        ws[f'H{r}'] = f'=IF(OR({nm}="",{SH_BASE}!${AC_REAL}${a}=""),"",{SH_BASE}!${AC_REAL}${a})'
        ws[f'I{r}'] = f'=IF(OR({nm}="",{SH_BASE}!${AC_RDATE}${a}=""),"",{SH_BASE}!${AC_RDATE}${a})'
        ws[f'J{r}'] = f'=IF(OR({nm}="",{SH_BASE}!${AC_DIFF}${a}=""),"",{SH_BASE}!${AC_DIFF}${a})'
        ws[f'K{r}'] = (f'=IF({nm}="","",IF(C{r}="个人户",IF(G{r}<0,"公司欠"&{SH_BASE}!${AC_OWNER}${a}&"的",IF(G{r}>0,"他手上公司的钱（备用金）","")),'
                       f'IF(C{r}="票据","在手的承兑汇票","")))')
    A1 = A0 + na - 1
    style_rows(ws, A0, A1, list('ABCDEFGHIJK'), auto=list('ABCDEFGHIJK'),
               fmts={c: MONEY for c in 'DEFGHJ'} | {'I': DATE}, aligns={'B': AL, 'K': AL, **{c: AR for c in 'DEFGHJ'}})
    for r in range(A0, A1 + 1):
        for c in 'ABCDEFGHIJK':
            ws[f'{c}{r}'].fill = FILL_NONE
    groups = [('银行＋现金（能动用的钱）', '"银行"', '"现金"'), ('承兑汇票', '"票据"', None), ('老板/负责人个人户（负＝公司欠）', '"个人户"', None)]
    for k, (lab, t1, t2) in enumerate(groups):
        r = A1 + 1 + k
        put(ws, f'B{r}', lab, F_TXTB, FILL_SUB, align=AL)
        ws.merge_cells(f'B{r}:C{r}')
        for c in 'DEFG':
            f = f'SUMIFS({c}${A0}:{c}${A1},$C${A0}:$C${A1},{t1})' + (f'+SUMIFS({c}${A0}:{c}${A1},$C${A0}:$C${A1},{t2})' if t2 else '')
            put(ws, f'{c}{r}', f'={f}', F_AUTOB, FILL_SUB, MONEY, AR)
        for c in 'AHIJK':
            put(ws, f'{c}{r}', None, F_AUTO, FILL_SUB)
    # ② 分段
    G0 = A1 + len(groups) + 3
    section(ws, G0, 'A', 'K', '② 分段（月报按周 · 周报按天 · 自定义长的按月；不含账户互转）', C_RPT)
    header(ws, G0 + 1, [('A', '段'), ('B', '起'), ('C', '止'), ('D', '收到的钱'), ('E', '付出的钱'), ('F', '净额\n（不含互转）'),
                        ('G', '银行＋现金\n段末余额'), ('H', '其中老板个人户\n代收付净额'), ('I', '银行现金跟个人户、\n票据互转（净）')],
           C_RPT, height=40)
    kind = f'IF($C$3="周报","日",IF({E}-{S}<=31,"周","月"))'
    bank_bal = lambda d: (f'SUMIFS({SH_BASE}!${AC_OPEN}${AC_R0}:${AC_OPEN}${AC_R1},{AC_TYPES_R},"银行")'
                          f'+SUMIFS({SH_BASE}!${AC_OPEN}${AC_R0}:${AC_OPEN}${AC_R1},{AC_TYPES_R},"现金")'
                          + ''.join(f'+SUMIFS({jr(J_NET)},{jr(J_ATYPE)},"{t}",{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<="&{d})'
                                    f'-SUMIFS({jr(J_NET)},{jr(J_TOTYPE)},"{t}",{jr(J_LINE)},"账户互转",{jr(J_DATE)},">="&{OPEN_DATE},{jr(J_DATE)},"<="&{d})'
                                    for t in ('银行', '现金')))
    s0 = G0 + 2
    for k in range(12):
        r = s0 + k
        if k == 0:
            ws[f'B{r}'] = f'={S}'
        else:
            ws[f'B{r}'] = f'=IF(OR(C{r - 1}="",C{r - 1}>={E}),"",C{r - 1}+1)'
        ws[f'A{r}'] = f'=IF(B{r}="","",{k + 1})'
        ws[f'C{r}'] = f'=IF(B{r}="","",MIN({E},IF({kind}="日",B{r},IF({kind}="周",B{r}+7-WEEKDAY(B{r},2),EOMONTH(B{r},0)))))'
        sd = f'{jr(J_DATE)},">="&B{r},{jr(J_DATE)},"<="&C{r}'
        ws[f'D{r}'] = f'=IF(B{r}="","",SUMIFS({jr(J_NET)},{jr(J_NET)},">0",{jr(J_LINE)},"<>账户互转",{sd}))'
        ws[f'E{r}'] = f'=IF(B{r}="","",-SUMIFS({jr(J_NET)},{jr(J_NET)},"<0",{jr(J_LINE)},"<>账户互转",{sd}))'
        ws[f'F{r}'] = f'=IF(B{r}="","",D{r}-E{r})'
        ws[f'G{r}'] = f'=IF(B{r}="","",ROUND({bank_bal(f"C{r}")},2))'
        ws[f'H{r}'] = f'=IF(B{r}="","",SUMIFS({jr(J_NET)},{jr(J_ATYPE)},"个人户",{jr(J_LINE)},"<>账户互转",{sd}))'
        prevg = f'$D${A1 + 1}' if k == 0 else f'G{r - 1}'
        ws[f'I{r}'] = f'=IF(B{r}="","",ROUND(G{r}-{prevg}-(F{r}-H{r}),2))'
    s1 = s0 + 11
    style_rows(ws, s0, s1, list('ABCDEFGHI'), auto=list('ABCDEFGHI'), fmts={'B': 'm/d', 'C': 'm/d', **{c: MONEY for c in 'DEFGHI'}},
               aligns={c: AR for c in 'DEFGHI'})
    r = s1 + 1
    put(ws, f'B{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'DEFHI':
        put(ws, f'{c}{r}', f'=SUM({c}{s0}:{c}{s1})', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'ACG':
        put(ws, f'{c}{r}', None, F_AUTO, FILL_TOT)
    # ③ 回款（按项目）、④ 付款前 15（按单位）
    HP, HU, HUK = 'X', 'Z', 'AA'       # 隐藏：项目回款；单位付款、排序键
    npj = PJ_R1 - PJ_R0 + 1
    for i in range(npj):
        rr = 1 + i
        pn = f'{SH_PROJ}!${PJ_NAME}${PJ_R0 + i}'
        ws[f'{HP}{rr}'] = f'=IF({pn}="",0,SUMIFS({jr(J_NET)},{jr(J_PJ)},{pn},{jr(J_LINE)},"应收账款",{DR}))'
        ws[f'{HP}{rr}'].font = F_HELP
    HPC = 'Y'
    counter(ws, HPC, 1, npj, lambda i: f'{HP}{1 + i}<>0')
    nu = UN_R1 - UN_R0 + 1
    for i in range(nu):
        rr = 1 + i
        un = f'{SH_UNIT}!${UN_NAME}${UN_R0 + i}'
        ws[f'{HU}{rr}'] = f'=IF({un}="",0,-SUMIFS({jr(J_NET)},{jr(J_UN)},{un},{jr(J_NET)},"<0",{jr(J_LINE)},"<>账户互转",{DR}))'
        ws[f'{HUK}{rr}'] = f'=IF({HU}{rr}>0,{HU}{rr}+({nu + 1}-{rr})/1000000,0)'
        ws[f'{HU}{rr}'].font = ws[f'{HUK}{rr}'].font = F_HELP
    YR = f'{jr(J_DATE)},">="&DATE(YEAR({E}),1,1),{jr(J_DATE)},"<="&{E}'
    P0 = r + 4
    section(ws, P0 - 2, 'A', 'K', '③ 本期收到的工程款（按项目）', C_AR)
    header(ws, P0 - 1, [('A', '序号'), ('B', '项目'), ('C', '本期收款'), ('D', '本年累计收款'), ('E', '应收余额\n（截止月末）'),
                        ('F', '累计确认'), ('G', '回款比例'), ('H', '甲方 / 总包')], C_AR, height=40)
    ws.merge_cells(f'H{P0 - 1}:K{P0 - 1}')
    for k in range(20):
        r = P0 + k
        ix = f'$W{r}'
        ws[ix] = f'={kth(k + 1, HPC, 1, npj)}'
        ws[ix].font = F_HELP
        pn = f'INDEX({PJ_NAMES},{ix})'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",{pn})'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX(${HP}$1:${HP}${npj},{ix}))'
        ws[f'D{r}'] = f'=IF({ix}=0,"",SUMIFS({jr(J_NET)},{jr(J_PJ)},B{r},{jr(J_LINE)},"应收账款",{YR}))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",INDEX({s_agg.ps(PS_AR_E)},{ix}))'
        ws[f'F{r}'] = f'=IF({ix}=0,"",INDEX({SH_AR}!$F${AR_R0}:$F${AR_R0 + npj - 1},{ix}))'
        ws[f'G{r}'] = f'=IF(OR({ix}=0,N(F{r})=0),"",INDEX({SH_AR}!$G${AR_R0}:$G${AR_R0 + npj - 1},{ix})/F{r})'
        ws[f'H{r}'] = f'=IF({ix}=0,"",INDEX({rng(SH_PROJ, PJ_CUS, PJ_R0, PJ_R1)},{ix})&"")'
        ws.merge_cells(f'H{r}:K{r}')
    style_rows(ws, P0, P0 + 19, list('ABCDEFGH'), auto=list('ABCDEFGH'), fmts={**{c: MONEY for c in 'CDEF'}, 'G': PCT},
               aligns={'B': AL, 'H': AL, **{c: AR for c in 'CDEF'}})
    r = P0 + 20
    put(ws, f'B{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'CD':
        put(ws, f'{c}{r}', f'=SUM({c}{P0}:{c}{r - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'AEFGH':
        put(ws, f'{c}{r}', None, fill_=FILL_TOT)
    U0 = r + 4
    section(ws, U0 - 2, 'A', 'K', '④ 本期付款最多的 15 家（单位 / 人，不含互转）', C_AR)
    header(ws, U0 - 1, [('A', '名次'), ('B', '单位 / 人'), ('C', '类型'), ('D', '本期付款'), ('E', '还欠他\n（截止月末）'), ('F', '说明')], C_AR,
           height=40)
    ws.merge_cells(f'F{U0 - 1}:K{U0 - 1}')
    for k in range(15):
        r = U0 + k
        ix = f'$W{r}'
        ws[ix] = f'=IFERROR(IF(LARGE(${HUK}$1:${HUK}${nu},{k + 1})<=0,0,MATCH(LARGE(${HUK}$1:${HUK}${nu},{k + 1}),${HUK}$1:${HUK}${nu},0)),0)'
        ws[ix].font = F_HELP
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX({UN_NAMES},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX({UN_TYPES_R},{ix})&"")'
        ws[f'D{r}'] = f'=IF({ix}=0,"",INDEX(${HU}$1:${HU}${nu},{ix}))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",INDEX({s_agg.bx(BX_AP_E)},{ix})+INDEX({s_agg.bx(BX_WG_E)},{ix}))'
        ws[f'F{r}'] = f'=IF({ix}=0,"",IF(N(E{r})<0,"多付了（他欠我们）",""))'
        ws.merge_cells(f'F{r}:K{r}')
    style_rows(ws, U0, U0 + 14, list('ABCDEF'), auto=list('ABCDEF'), fmts={'D': MONEY, 'E': MONEY}, aligns={'B': AL, 'D': AR, 'E': AR, 'F': AL})
    # ⑤ 收支分类
    C0 = U0 + 15 + 3
    section(ws, C0 - 2, 'A', 'K', '⑤ 本期收支分类（按收支类别；账户互转单列，不算收支）', C_RPT)
    header(ws, C0 - 1, [('A', '序号'), ('B', '收支类别'), ('C', '报表项目'), ('D', '收入\n（银行现金票据）'), ('E', '支出\n（银行现金票据）'),
                        ('F', '个人户\n代公司收'), ('G', '个人户\n代公司付'), ('H', '本期净额'), ('I', '笔数'), ('J', '本年累计\n净额'), ('K', '说明')],
           C_RPT, height=40)
    nc = CT_R1 - CT_R0 + 1
    for i in range(nc):
        r, c_ = C0 + i, CT_R0 + i
        cat = f'$B{r}'
        ws[f'A{r}'] = f'=IF({cat}="","",{i + 1})'
        ws[f'B{r}'] = f'={SH_BASE}!${CT_NAME}${c_}&""'
        ws[f'C{r}'] = f'={SH_BASE}!${CT_LINE}${c_}&""'
        ws[f'D{r}'] = (f'=IF({cat}="","",SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_NET)},">0",{DR})'
                       f'-SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_ATYPE)},"个人户",{jr(J_NET)},">0",{DR}))')
        ws[f'E{r}'] = (f'=IF({cat}="","",-SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_NET)},"<0",{DR})'
                       f'+SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_ATYPE)},"个人户",{jr(J_NET)},"<0",{DR}))')
        ws[f'F{r}'] = f'=IF({cat}="","",SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_ATYPE)},"个人户",{jr(J_NET)},">0",{DR}))'
        ws[f'G{r}'] = f'=IF({cat}="","",-SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{jr(J_ATYPE)},"个人户",{jr(J_NET)},"<0",{DR}))'
        ws[f'H{r}'] = f'=IF({cat}="","",D{r}-E{r}+F{r}-G{r})'
        ws[f'I{r}'] = f'=IF({cat}="","",COUNTIFS({jr(J_CT)},{cat},{DR}))'
        ws[f'J{r}'] = f'=IF({cat}="","",SUMIFS({jr(J_NET)},{jr(J_CT)},{cat},{YR}))'
        ws[f'K{r}'] = f'=IF({cat}="账户互转","自己账户之间倒钱，不算收支","")'
    C1 = C0 + nc - 1
    style_rows(ws, C0, C1, list('ABCDEFGHIJK'), auto=list('ABCDEFGHIJK'), fmts={**{c: MONEY for c in 'DEFGHJ'}, 'I': INT},
               aligns={'B': AL, 'C': AL, 'K': AL, **{c: AR for c in 'DEFGHJ'}})
    for r in range(C0, C1 + 1):
        for c in 'ABCDEFGHIJK':
            ws[f'{c}{r}'].fill = FILL_NONE
    ws.conditional_formatting.add(f'A{C0}:K{C1}', FormulaRule(formula=[f'AND($B{C0}<>"",$I{C0}=0)'], font=GREY))
    ws.conditional_formatting.add(f'A{C0}:K{C1}', FormulaRule(formula=[f'$B{C0}="账户互转"'], fill=fill('FFF2F2F2')))
    r = C1 + 1
    put(ws, f'B{r}', '没认出类别的', F_TXT, align=AL)
    put(ws, f'H{r}', f'=SUMIFS({jr(J_NET)},{jr(J_LINE)},"待分类",{DR})', F_AUTO, fmt=MONEY, align=AR)
    put(ws, f'I{r}', f'=COUNTIFS({jr(J_LINE)},"待分类",{DR})', F_AUTO, fmt=INT, align=AC)
    put(ws, f'K{r}', '去【资金流水】看 ✗ 的行', F_NOTE, align=AL)
    for c in 'ACDEFGJ':
        put(ws, f'{c}{r}', None)
    r += 1
    put(ws, f'B{r}', '合计（不含账户互转）', F_TXTB, FILL_TOT, align=AL)
    ws.merge_cells(f'B{r}:C{r}')
    for c in 'DEFGHJ':
        put(ws, f'{c}{r}', f'=SUM({c}{C0}:{c}{r - 1})-SUMIFS({c}{C0}:{c}{C1},$B{C0}:$B{C1},"账户互转")', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'I{r}', f'=SUM(I{C0}:I{r - 1})-SUMIFS(I{C0}:I{C1},$B{C0}:$B{C1},"账户互转")', F_AUTOB, FILL_TOT, INT, AC)
    for c in 'AK':
        put(ws, f'{c}{r}', None, fill_=FILL_TOT)
    hide(ws, 'W', HP, HPC, HU, HUK)
    ws.freeze_panes = 'A6'
    print_setup(ws, None, landscape=False)
    return ws


# ═══════════════════════════ 应收账款总表 ═══════════════════════════
AR_R0 = 6


def build_ar(wb, ctx):
    ws = _sheet(wb, SH_AR, C_AR, '应 收 账 款 总 表（每个项目：确认了多少、收了多少、还欠多少、欠了多久）',
                '💡 全自动，到首页选的截止日为止。应收金额＝甲方/总包确认的产值（【收入确认】＋期初），已收＝收到的工程款＋总包代发工资等抵账。'
                '质保金没到期的不算「可催收」。账龄按「先确认的先收回」算：还欠的钱是最近几次确认的那些。'
                '开票未回款＝开了票还没收到钱的（要重点催）。点项目名跳到【项目账】，在那边黄格子选这个项目就能看全部明细。',
                'Y', {'A': 5, 'B': 14, 'C': 24, 'D': 13, 'E': 13, 'F': 14, 'G': 14, 'H': 8, 'I': 14, 'J': 13, 'K': 13, 'L': 13,
                      'M': 13, 'N': 11, 'O': 13, 'P': 11, 'Q': 8, 'R': 12, 'S': 12, 'T': 12, 'U': 12, 'V': 10, 'W': 13, 'X': 11, 'Y': 26})
    put(ws, 'B3', f'="截止 "&TEXT({AX_E},"yyyy年m月d日")', F_KPI_L, align=AL, border=False)
    ws.merge_cells('B3:D3')
    heads = [('A', '序号'), ('B', '项目'), ('C', '甲方 / 总包'), ('D', '合同额'), ('E', '总价\n（结算/合同＋变更）'), ('F', '应收金额\n（累计确认产值）'),
             ('G', '已收金额\n（含抵账）'), ('H', '回款\n比例'), ('I', '未收金额\n（应收余额）'), ('J', '已开票'), ('K', '未开票\n（确认−开票）'),
             ('L', '开票未回款'), ('M', '质保金'), ('N', '质保到期'), ('O', '可催收\n（不含未到期质保）'), ('P', '最后回款日'),
             ('Q', '多久\n没回款(天)'), ('R', '3个月内'), ('S', '3～6个月'), ('T', '6～12个月'), ('U', '1年以上\n（含建账前）'), ('V', '状态'),
             ('W', '还没确认的\n（总价−确认）'), ('X', '确认进度'), ('Y', '提示')]
    header(ws, AR_R0 - 1, heads, C_AR, height=40)
    ws.row_dimensions[4].height = 20
    npj = PJ_R1 - PJ_R0 + 1
    Ed = AX_E
    for i in range(npj):
        r, p, k = AR_R0 + i, PJ_R0 + i, PS_R0 + i
        P = lambda col: f'{SH_PROJ}!${col}${p}'
        pc = lambda m, per: f'{SH_PS}!${PS_COL[(m, per)]}${k}'
        nm = f'$B{r}'
        ws[f'A{r}'] = f'=IF({nm}="","",{i + 1})'
        ws[f'B{r}'] = f'=IF({P(PJ_NAME)}="","",HYPERLINK("#\'{SH_PL}\'!C3",{P(PJ_NAME)}&""))'
        ws[f'C{r}'] = f'={P(PJ_CUS)}&""'
        ws[f'D{r}'] = f'=IF({nm}="","",N({P(PJ_AMT)}))'
        ws[f'E{r}'] = f'=IF({nm}="","",N({P(PJ_TOTAL)}))'
        o = lambda col: f'SUMIFS({opr(col)},{opr(OPJ_PJ)},{nm})'
        ws[f'F{r}'] = f'=IF({nm}="","",{o(OPJ_REV)}+{pc("确认收入", "前")}+{pc("确认收入", "本")})'
        ws[f'G{r}'] = f'=IF({nm}="","",{o(OPJ_REC)}+{pc("现金回款", "前")}+{pc("现金回款", "本")}+{pc("抵账", "前")}+{pc("抵账", "本")})'
        ws[f'H{r}'] = f'=IF(OR({nm}="",N(F{r})=0),"",G{r}/F{r})'
        ws[f'I{r}'] = f'=IF({nm}="","",{SH_PS}!${PS_AR_E}${k})'
        ws[f'J{r}'] = f'=IF({nm}="","",{o(OPJ_INV)}+{pc("开票", "前")}+{pc("开票", "本")})'
        ws[f'K{r}'] = f'=IF({nm}="","",F{r}-J{r})'
        ws[f'L{r}'] = f'=IF({nm}="","",MAX(J{r}-G{r},0))'
        ws[f'M{r}'] = f'=IF({nm}="","",ROUND(E{r}*N({P(PJ_RET)}),2))'
        ws[f'N{r}'] = f'=IF(OR({nm}="",NOT(ISNUMBER({P(PJ_RETD)}))),"",{P(PJ_RETD)})'
        ws[f'O{r}'] = f'=IF({nm}="","",MAX(0,I{r}-IF(AND(ISNUMBER(N{r}),N{r}<={Ed}),0,MIN(MAX(I{r},0),M{r}))))'
        last_j = f'SUMPRODUCT(MAX(({jr(J_PJ)}={nm})*({jr(J_LINE)}="应收账款")*({jr(J_NET)}>0)*({jr(J_DN)}<={Ed})*{jr(J_DN)}))'
        last_o = (f'SUMPRODUCT(MAX(({ofr(OF_PJ)}={nm})*(({ofr(OF_TYPE)}="总包代发工资")+({ofr(OF_TYPE)}="总包代付材料分包款"))'
                  f'*({ofr(OF_DN)}<={Ed})*{ofr(OF_DN)}))')
        ws[f'P{r}'] = f'=IF({nm}="","",IF(MAX({last_j},{last_o})=0,"",MAX({last_j},{last_o})))'
        ws[f'Q{r}'] = f'=IF(OR({nm}="",N(I{r})<=0),"",IF(P{r}="","建账后没收过",{Ed}-P{r}))'
        win = lambda days: (f'SUMIFS({rvr(RV_AMT)},{rvr(RV_PJ)},{nm},{rvr(RV_AMT)},">0",{rvr(RV_DATE)},">"&({Ed}-{days}),{rvr(RV_DATE)},"<="&{Ed})')
        B = f'MAX(I{r},0)'
        m90, m180, m365 = (f'MIN({B},{win(d)})' for d in (91, 183, 365))
        ws[f'R{r}'] = f'=IF({nm}="","",{m90})'
        ws[f'S{r}'] = f'=IF({nm}="","",{m180}-{m90})'
        ws[f'T{r}'] = f'=IF({nm}="","",{m365}-{m180})'
        ws[f'U{r}'] = f'=IF({nm}="","",{B}-{m365})'
        ws[f'V{r}'] = f'={P(PJ_STAT)}&""'
        ws[f'W{r}'] = f'=IF(OR({nm}="",N(E{r})=0),"",E{r}-F{r})'
        ws[f'X{r}'] = f'=IF(OR({nm}="",N(E{r})=0),"",F{r}/E{r})'
        ws[f'Y{r}'] = (f'=IF({nm}="","",IF(N(I{r})<0,"预收了（收的比确认的多，记得补确认产值）",'
                       f'IF(AND(ISNUMBER(N{r}),N{r}<={Ed},N(I{r})>0,N(M{r})>0),"质保金已到期，可以催",'
                       f'IF(N(L{r})>0,"开了票没收到钱",IF(AND(ISNUMBER(Q{r}),N(Q{r})>180),"半年多没回款","")))))')
    R1 = AR_R0 + npj - 1
    cols = [CL(i) for i in range(1, 26)]
    style_rows(ws, AR_R0, R1, cols, auto=cols,
               fmts={**{c: MONEY for c in 'DEFGIJKLMORSTUW'}, 'H': PCT, 'X': PCT, 'N': DATE, 'P': DATE, 'Q': INT},
               aligns={'B': AL, 'C': AL, 'Y': AL, **{c: AR for c in 'DEFGIJKLMORSTUW'}})
    for r in range(AR_R0, R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'B{r}'].font = Font(name=YH, sz=10, color='FF0563C1', underline='single')
    ws.conditional_formatting.add(f'Y{AR_R0}:Y{R1}', FormulaRule(formula=[f'$Y{AR_R0}<>""'], font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'U{AR_R0}:U{R1}', FormulaRule(formula=[f'N($U{AR_R0})>0'], fill=FILL_WARN))
    # 合计（第 4 行）
    put(ws, 'B4', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'DEFGIJKLMORSTUW':
        put(ws, f'{c}4', f'=SUM({c}{AR_R0}:{c}{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'H4', f'=IF(N(F4)=0,"",G4/F4)', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, 'X4', f'=IF(N(E4)=0,"",F4/E4)', F_AUTOB, FILL_TOT, PCT, AC)
    for c in 'ACNPQVY':
        put(ws, f'{c}4', None, fill_=FILL_TOT)
    ws.freeze_panes = f'C{AR_R0}'
    ws.auto_filter.ref = f'A{AR_R0 - 1}:Y{R1}'
    return ws


# ═══════════════════════════ 应付账款总表 ＋ 单位对账单 ═══════════════════════════
def build_aps(wb, ctx):
    last = CL(11 + APS_PJ)
    ws = _sheet(wb, SH_APS, C_AR, '应 付 账 款 总 表（材料商 · 分包 · 机械运输：欠谁多少、各项目多少）＋ 对账单',
                '💡 黄格子选类型：「材料供应商」＝你原来的应付账款总表，「分包」＝应付分包商账款，「机械运输」＝机械费汇总；「全部」＝都列出来。'
                '应付＝【应付登记】＋【期初余额】②，已付＝【资金流水】里付给他的（类别「付应付款」）＋总包代付。后面几列是这家在各项目的应付。'
                '最下面「对账单」：选一家，列出他所有的应付和付款，跟他对账用。',
                last, {'A': 5, 'B': 18, 'C': 10, 'D': 13, 'E': 13, 'F': 13, 'G': 8, 'H': 13, 'I': 13, 'J': 11, 'K': 12,
                       **{CL(12 + i): 12 for i in range(APS_PJ)}})
    _lbl(ws, 'B3', '类型')
    put(ws, 'C3', '全部', F_SEL, FILL_SEL, align=AC)
    dv_list(ws, 'C3', '"全部,材料供应商,分包,机械运输,其他"')
    put(ws, 'D3', f'="截止 "&TEXT({AX_E},"yyyy年m月d日")', F_KPI_L, align=AL, border=False)
    ws.merge_cells('D3:F3')
    T = '$C$3'
    tc = f'IF({T}="全部","*",{T})'
    nu = UN_R1 - UN_R0 + 1
    npj = PJ_R1 - PJ_R0 + 1
    HU, HP = 'AZ', 'BA'          # 隐藏计数：单位、项目
    AP_T = ('材料供应商', '分包', '机械运输', '其他')
    def ucond(i):
        b = BX_R0 + i
        t = f'{SH_BALX}!${BX_TYPE}${b}'
        typ_ok = f'OR({t}="{AP_T[0]}",{t}="{AP_T[1]}",{t}="{AP_T[2]}",{t}="{AP_T[3]}")'
        return f'AND({typ_ok},OR({T}="全部",{t}={T}),OR({SH_BALX}!${BX_APAMT}${b}<>0,{SH_BALX}!${BX_PAID}${b}<>0))'
    counter(ws, HU, 1, nu, ucond)
    def pcond(i):
        pn = f'{SH_PROJ}!${PJ_NAME}${PJ_R0 + i}'
        amt = (f'SUMIFS({oapr(OAP_AMT)},{oapr(OAP_PJ)},{pn},{oapr(OAP_TYPE)},{tc})'
               f'+SUMIFS({apr(AP_AMT)},{apr(AP_PJ)},{pn},{apr(AP_UTYPE)},{tc},{apr(AP_YM)},">="&{AX_OYM},{apr(AP_YM)},"<="&{AX_YM1})')
        return f'AND({pn}<>"",{amt}<>0)'
    counter(ws, HP, 1, npj, pcond)
    H0 = 6
    R0 = H0 + 1
    heads = [('A', '序号'), ('B', '单位'), ('C', '类型'), ('D', '应付总额'), ('E', '已付总额'), ('F', '剩余未付'), ('G', '付款\n比例'),
             ('H', '已收票'), ('I', '未收票\n（应付−收票）'), ('J', '最后付款日'), ('K', '没分项目\n或其他项目')]
    header(ws, H0, heads, C_AR, height=40)
    for j in range(APS_PJ):
        c = CL(12 + j)
        put(ws, f'{c}{H0}', f'=IFERROR(INDEX({PJ_NAMES},MATCH({j + 1},${HP}$1:${HP}${npj},0)),"")', F_HDR, fill(C_AR), align=ACW)
    for k in range(APS_ROWS):
        r = R0 + k
        ix = f'$AY{r}'
        ws[ix] = f'={kth(k + 1, HU, 1, nu)}'
        ws[ix].font = F_HELP
        b = lambda col: f'INDEX({s_agg.bx(col)},{ix})'
        nm = f'$B{r}'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX({UN_NAMES},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX({UN_TYPES_R},{ix}))'
        ws[f'D{r}'] = f'=IF({ix}=0,"",{b(BX_APAMT)})'
        ws[f'E{r}'] = f'=IF({ix}=0,"",{b(BX_PAID)})'
        ws[f'F{r}'] = f'=IF({ix}=0,"",{b(BX_AP_E)})'
        ws[f'G{r}'] = f'=IF(OR({ix}=0,N(D{r})=0),"",E{r}/D{r})'
        ws[f'H{r}'] = f'=IF({ix}=0,"",{b(BX_INVIN)})'
        ws[f'I{r}'] = f'=IF({ix}=0,"",D{r}-H{r})'
        lp = f'SUMPRODUCT(MAX(({jr(J_UN)}={nm})*({jr(J_NET)}<0)*({jr(J_DN)}<={AX_E})*{jr(J_DN)}))'
        ws[f'J{r}'] = f'=IF({ix}=0,"",IF({lp}=0,"",{lp}))'
        pcols = []
        for j in range(APS_PJ):
            c = CL(12 + j)
            h = f'{c}${H0}'
            ws[f'{c}{r}'] = (f'=IF(OR({ix}=0,{h}=""),"",SUMIFS({oapr(OAP_AMT)},{oapr(OAP_UNIT)},{nm},{oapr(OAP_PJ)},{h})'
                             f'+SUMIFS({apr(AP_AMT)},{apr(AP_UNIT)},{nm},{apr(AP_PJ)},{h},{apr(AP_YM)},">="&{AX_OYM},{apr(AP_YM)},"<="&{AX_YM1}))')
            pcols.append(f'N({c}{r})')
        ws[f'K{r}'] = f'=IF({ix}=0,"",D{r}-({"+".join(pcols)}))'
    R1 = R0 + APS_ROWS - 1
    cols = [CL(i) for i in range(1, 12 + APS_PJ)]
    style_rows(ws, R0, R1, cols, auto=cols, fmts={**{c: MONEY for c in cols if c not in 'ABCGJ'}, 'G': PCT, 'J': DATE},
               aligns={'B': AL, **{c: AR for c in cols if c not in 'ABCGJ'}})
    for r in range(R0, R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    ws.conditional_formatting.add(f'F{R0}:F{R1}', FormulaRule(formula=[f'N($F{R0})<0'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'F{R0}:F{R1}', FormulaRule(formula=[f'N($F{R0})>0'], font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    put(ws, 'B4', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in cols:
        if c in 'ABCGJ':
            continue
        put(ws, f'{c}4', f'=SUM({c}{R0}:{c}{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'G4', '=IF(N(D4)=0,"",E4/D4)', F_AUTOB, FILL_TOT, PCT, AC)
    for c in 'ACJ':
        put(ws, f'{c}4', None, fill_=FILL_TOT)
    put(ws, 'G3', f'=IF({cnt(HU, 1, nu)}>{APS_ROWS},"⚠ 有"&{cnt(HU, 1, nu)}&"家，只列前{APS_ROWS}家","")&IF({cnt(HP, 1, npj)}>{APS_PJ},'
                  f'"  ⚠ 有"&{cnt(HP, 1, npj)}&"个项目，前{APS_PJ}个单列，其余在「没分项目或其他项目」","")', F_RED, border=False)
    ws.merge_cells(f'G3:{last}3')
    hide(ws, 'AY', HU, HP)
    # ── 对账单 ──
    S0 = R1 + 3
    section(ws, S0, 'A', last, '对 账 单（选一家单位：他所有的应付和付款）', C_CHK)
    _lbl(ws, f'B{S0 + 1}', '单位')
    first = next((u['name'] for u in ctx['units'] if u['type'] == '材料供应商'), '')
    put(ws, f'C{S0 + 1}', first, F_SEL, FILL_SEL, align=AC)
    ws.merge_cells(f'C{S0 + 1}:D{S0 + 1}')
    dv_list(ws, f'C{S0 + 1}', f'={UN_NAMES}', '选一家单位', stop=False)
    U = f'$C${S0 + 1}'
    bxu = lambda col: f'IFERROR(INDEX({s_agg.bx(col)},MATCH({U},{UN_NAMES},0)),0)'
    for c, t, f in (('E', '应付合计', bxu(BX_APAMT)), ('G', '已付合计', bxu(BX_PAID)), ('I', '还欠', bxu(BX_AP_E))):
        _lbl(ws, f'{c}{S0 + 1}', t)
        _val(ws, f'{CL(CI(c) + 1)}{S0 + 1}', f'={f}', font=F_KPI_V if t == '还欠' else F_AUTOB)
    _lbl(ws, f'K{S0 + 1}', '已收票')
    _val(ws, f'L{S0 + 1}', f'={bxu(BX_INVIN)}')
    h = S0 + 3
    header(ws, h, [('A', '序号'), ('B', '日期'), ('C', '项目'), ('D', '内容'), ('E', '应付金额')], C_AR)
    header(ws, h, [('G', '序号'), ('H', '日期'), ('I', '付款账户'), ('J', '摘要'), ('K', '付款金额')], C_CASH)
    put(ws, f'B{h - 1}', '应付（期初 ＋ 应付登记）', F_TXTB, border=False)
    put(ws, f'H{h - 1}', '付款（资金流水 ＋ 总包代付）', F_TXTB, border=False)
    # 应付清单：期初行（200）＋ 应付登记行（1000）；付款：流水（2000）＋ 代发抵账（300）
    HA, HB, HC_, HD = 'BB', 'BC', 'BD', 'BE'
    no = OAP_R1 - OAP_R0 + 1
    na_ = AP_R1 - AP_R0 + 1
    nj = J_R1 - J_R0 + 1
    nf = OF_R1 - OF_R0 + 1
    counter(ws, HA, 1, no, lambda i: f'AND({U}<>"",{SH_OPEN}!${OAP_UNIT}${OAP_R0 + i}={U})')
    counter(ws, HB, 1, na_, lambda i: (f'({U}<>"")*({SH_AP}!${AP_UNIT}${AP_R0 + i}={U})*ISNUMBER({SH_AP}!${AP_YM}${AP_R0 + i})'
                                        f'*({SH_AP}!${AP_YM}${AP_R0 + i}>={AX_OYM})*({SH_AP}!${AP_YM}${AP_R0 + i}<={AX_YM1})'))
    NCC = skey(ws, HC_, 2, nj, lambda i: (f'AND({U}<>"",{jc(J_UN, J_R0 + i)}={U},ISNUMBER({jc(J_DATE, J_R0 + i)}),ISNUMBER({jc(J_YM, J_R0 + i)}),'
                                           f'{jc(J_YM, J_R0 + i)}>={AX_OYM},{jc(J_YM, J_R0 + i)}<={AX_YM1},'
                                           f'OR({jc(J_LINE, J_R0 + i)}="应付材料款",{jc(J_LINE, J_R0 + i)}="应付分包款",'
                                           f'{jc(J_LINE, J_R0 + i)}="应付机械运输费",{jc(J_LINE, J_R0 + i)}="应付其他款"))'),
               lambda i: f'INT({jc(J_DATE, J_R0 + i)})')
    counter(ws, HD, 1, nf, lambda i: (f'({U}<>"")*({SH_OFF}!${OF_WHO}${OF_R0 + i}={U})*({SH_OFF}!${OF_TYPE}${OF_R0 + i}="总包代付材料分包款")'
                                       f'*ISNUMBER({SH_OFF}!${OF_YM}${OF_R0 + i})*({SH_OFF}!${OF_YM}${OF_R0 + i}<={AX_YM1})'))
    ca, cb, cc, cd = cnt(HA, 1, no), cnt(HB, 1, na_), NCC, cnt(HD, 1, nf)
    for k in range(STMT_CAP):
        r = h + 1 + k
        # 左：先期初，再应付登记
        ka, kb = f'{k + 1}', f'{k + 1}-{ca}'
        ia = f'IF({k + 1}<={ca},{kth(ka, HA, 1, no)},0)'
        ib = f'IF(AND({k + 1}>{ca},{k + 1}<={ca}+{cb}),{kth(kb, HB, 1, na_)},0)'
        ws[f'AW{r}'], ws[f'AX{r}'] = f'={ia}', f'={ib}'
        ws[f'AW{r}'].font = ws[f'AX{r}'].font = F_HELP
        A_, B_ = f'$AW{r}', f'$AX{r}'
        oa = lambda col: f'INDEX({oapr(col)},{A_})'
        ab = lambda col: f'INDEX({apr(col)},{B_})'
        ws[f'A{r}'] = f'=IF({A_}+{B_}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({A_}>0,"期初",IF({B_}>0,{ab(AP_DATE)},""))'
        ws[f'C{r}'] = f'=IF({A_}>0,{oa(OAP_PJ)}&"",IF({B_}>0,{ab(AP_PJ)}&"",""))'
        ws[f'D{r}'] = (f'=IF({A_}>0,"建账前累计（已付 "&TEXT(N({oa(OAP_PAID)}),"#,##0.00")&"）"&IF({oa(OAP_NOTE)}="","","；"&{oa(OAP_NOTE)}),'
                       f'IF({B_}>0,{ab(AP_MEMO)}&"",""))')
        ws[f'E{r}'] = f'=IF({A_}>0,N({oa(OAP_AMT)}),IF({B_}>0,N({ab(AP_AMT)}),""))'
        # 右：先流水付款，再总包代付；期初已付放第一行
        kc = f'{k}'
        ic = f'IF(AND({k}>=1,{k}<={cc}),{ksorted(kc, HC_, 2, nj, NCC)},0)'
        id_ = f'IF(AND({k}>{cc},{k}<={cc}+{cd}),{kth(f"{k}-{cc}", HD, 1, nf)},0)'
        ws[f'AU{r}'], ws[f'AV{r}'] = f'={ic}', f'={id_}'
        ws[f'AU{r}'].font = ws[f'AV{r}'].font = F_HELP
        C_, D_ = f'$AU{r}', f'$AV{r}'
        jj = lambda col: f'INDEX({jr(col)},{C_})'
        ff = lambda col: f'INDEX({ofr(col)},{D_})'
        if k == 0:
            ws[f'G{r}'] = '=1'
            ws[f'H{r}'] = '="期初"'
            ws[f'I{r}'] = '=""'
            ws[f'J{r}'] = '="建账前累计已付（【期初余额】②）"'
            ws[f'K{r}'] = f'=SUMIFS({oapr(OAP_PAID)},{oapr(OAP_UNIT)},{U})'
        else:
            ws[f'G{r}'] = f'=IF({C_}+{D_}=0,"",{k + 1})'
            ws[f'H{r}'] = f'=IF({C_}>0,{jj(J_DATE)},IF({D_}>0,{ff(OF_DATE)},""))'
            ws[f'I{r}'] = f'=IF({C_}>0,{jj(J_ACC)}&"",IF({D_}>0,"总包代付",""))'
            ws[f'J{r}'] = f'=IF({C_}>0,{jj(J_MEMO)}&"",IF({D_}>0,{ff(OF_PJ)}&" "&{ff(OF_NOTE)},""))'
            ws[f'K{r}'] = f'=IF({C_}>0,-{jj(J_NET)},IF({D_}>0,N({ff(OF_AMT)}),""))'
    r1 = h + STMT_CAP
    style_rows(ws, h + 1, r1, list('ABCDEGHIJK'), auto=list('ABCDEGHIJK'), fmts={'B': DATE, 'H': DATE, 'E': MONEY, 'K': MONEY},
               aligns={'C': AL, 'D': AL, 'I': AL, 'J': AL, 'E': AR, 'K': AR})
    for r in range(h + 1, r1 + 1):
        for c in 'ABCDEGHIJK':
            ws[f'{c}{r}'].fill = FILL_NONE
    hide(ws, 'AU', 'AV', 'AW', 'AX', HA, HB, HC_, HD)
    put(ws, f'M{S0 + 1}', f'=IF({ca}+{cb}>{STMT_CAP},"⚠ 应付超过{STMT_CAP}笔，只列前{STMT_CAP}笔","")&IF({cc}+{cd}>{STMT_CAP - 1},"⚠ 付款超过{STMT_CAP - 1}笔，只列前面的","")',
        F_RED, border=False)
    ws.freeze_panes = f'C{R0}'
    return ws


# ═══════════════════════════ 个人往来（老板、负责人的个人户） ═══════════════════════════
def build_per(wb, ctx):
    ws = _sheet(wb, SH_PER, C_CASH, '个 人 往 来（老板、负责人替公司收付的钱 · 公司欠他多少）',
                '💡 上面：每个人的个人户一行——他替公司付了多少、替公司收了多少、公司转给他多少（还他钱、报销、备用金）、他转给公司多少（借给公司），'
                '余额是负数＝公司欠他，正数＝他手上还有公司的钱（备用金没花完）。另外公司欠他的工资单独算。'
                '下面：选一个人，列出跟他有关的每一笔（跟你原来的「王总个人收支明细」一个意思）。右上是保证金、押金台账。',
                'M', {'A': 6, 'B': 11, 'C': 34, 'D': 13, 'E': 13, 'F': 13, 'G': 13, 'H': 13, 'I': 13, 'J': 12, 'K': 12, 'L': 14, 'M': 14})
    put(ws, 'B3', f'="截止 "&TEXT({AX_E},"yyyy年m月d日")', F_KPI_L, align=AL, border=False)
    ws.merge_cells('B3:D3')
    H0 = 5
    header(ws, H0, [('A', '序号'), ('B', '个人户'), ('C', '是谁的'), ('D', '期初余额\n（建账日）'), ('E', '他替公司付'), ('F', '他替公司收'),
                    ('G', '公司转给他\n（还款/报销/备用金）'), ('H', '他转给公司\n（借给公司）'), ('I', '个人户余额'), ('J', '公司欠他工资'),
                    ('K', '合计：公司欠他'), ('L', '说明')], C_CASH, height=44)
    ws.merge_cells(f'L{H0}:M{H0}')
    HC = 'Z'
    na = AC_R1 - AC_R0 + 1
    counter(ws, HC, 1, na, lambda i: f'{SH_BASE}!${AC_TYPE}${AC_R0 + i}="个人户"')
    rj = f'{jr(J_YM)},">="&{AX_OYM},{jr(J_YM)},"<="&{AX_YM1}'
    NP = 8
    for k in range(NP):
        r = H0 + 1 + k
        ix = f'$Y{r}'
        ws[ix] = f'={kth(k + 1, HC, 1, na)}'
        ws[ix].font = F_HELP
        nm = f'$B{r}'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX({AC_NAMES},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX({AC_OWNERS},{ix})&"")'
        ws[f'D{r}'] = f'=IF({ix}=0,"",N(INDEX({AC_OPENS},{ix})))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",-SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_LINE)},"<>账户互转",{jr(J_NET)},"<0",{rj}))'
        ws[f'F{r}'] = f'=IF({ix}=0,"",SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_LINE)},"<>账户互转",{jr(J_NET)},">0",{rj}))'
        ws[f'G{r}'] = (f'=IF({ix}=0,"",SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},">0",{rj})'
                       f'-SUMIFS({jr(J_NET)},{jr(J_TO)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},"<0",{rj}))')
        ws[f'H{r}'] = (f'=IF({ix}=0,"",-SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},"<0",{rj})'
                       f'+SUMIFS({jr(J_NET)},{jr(J_TO)},{nm},{jr(J_LINE)},"账户互转",{jr(J_NET)},">0",{rj}))')
        ws[f'I{r}'] = f'=IF({ix}=0,"",ROUND(D{r}-E{r}+F{r}+G{r}-H{r},2))'
        ws[f'J{r}'] = f'=IF({ix}=0,"",IFERROR(INDEX({s_agg.bx(BX_WG_E)},MATCH(C{r},{UN_NAMES},0)),0))'
        ws[f'K{r}'] = f'=IF({ix}=0,"",-I{r}+J{r})'
        ws[f'L{r}'] = (f'=IF({ix}=0,"",IF(K{r}>0,"公司一共欠"&C{r}&" "&TEXT(K{r},"#,##0.00"),IF(K{r}<0,C{r}&"手上还有公司的 "&TEXT(-K{r},"#,##0.00"),"两清")))')
        ws.merge_cells(f'L{r}:M{r}')
    R1 = H0 + NP
    style_rows(ws, H0 + 1, R1, list('ABCDEFGHIJKL'), auto=list('ABCDEFGHIJKL'), fmts={c: MONEY for c in 'DEFGHIJK'},
               aligns={'L': AL, 'C': AC, **{c: AR for c in 'DEFGHIJK'}}, bold=['K'])
    for r in range(H0 + 1, R1 + 1):
        for c in 'ABCDEFGHIJKL':
            ws[f'{c}{r}'].fill = FILL_NONE
    r = R1 + 1
    put(ws, f'B{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'DEFGHIJK':
        put(ws, f'{c}{r}', f'=SUM({c}{H0 + 1}:{c}{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'ACL':
        put(ws, f'{c}{r}', None, fill_=FILL_TOT)
    ws.merge_cells(f'L{r}:M{r}')
    # ── 选一个人看明细 ──
    D0 = R1 + 4
    section(ws, D0, 'A', 'M', '选一个人看明细（跟他有关的每一笔：他个人户记的、别的账户跟他互转的、发给他的工资）', C_CASH)
    _lbl(ws, f'B{D0 + 1}', '个人户')
    first = next((a[0] for a in ctx['accounts'] if a[1] == '个人户'), '')
    put(ws, f'C{D0 + 1}', first, F_SEL, FILL_SEL, align=AC)
    dv_list(ws, f'C{D0 + 1}', f'={AC_NAMES}', '选他的个人户')
    A = f'$C${D0 + 1}'
    OW = f'$E${D0 + 1}'
    _lbl(ws, f'D{D0 + 1}', '是谁的')
    _val(ws, OW.replace('$', ''), f'=IFERROR(INDEX({AC_OWNERS},MATCH({A},{AC_NAMES},0))&"","")', None, F_AUTOB)
    h = D0 + 3
    header(ws, h, [('A', '序号'), ('B', '日期'), ('C', '摘要'), ('D', '公司转给他'), ('E', '他转给公司'), ('F', '他替公司付'), ('G', '他替公司收'),
                   ('H', '发给他的工资'), ('I', '个人户余额\n（负＝公司欠他）'), ('J', '收支类别'), ('K', '项目'), ('L', '记在哪个账户'), ('M', '备注')],
           C_CASH, height=40)
    n = J_R1 - J_R0 + 1
    HD = 'AA'
    NDC = skey(ws, HD, J_R0, n, lambda i: (
        f'AND(ISNUMBER({jc(J_DATE, J_R0 + i)}),ISNUMBER({jc(J_YM, J_R0 + i)}),{jc(J_YM, J_R0 + i)}>={AX_OYM},{jc(J_YM, J_R0 + i)}<={AX_YM1},'
        f'OR({jc(J_ACC, J_R0 + i)}={A},AND({jc(J_TO, J_R0 + i)}={A},{jc(J_LINE, J_R0 + i)}="账户互转"),'
        f'AND({OW}<>"",{jc(J_UN, J_R0 + i)}={OW},{jc(J_LINE, J_R0 + i)}="应付职工薪酬")))'),
        lambda i: f'INT({jc(J_DATE, J_R0 + i)})')
    put(ws, f'G{D0 + 1}', f'=IF({NDC}>{PER_CAP},"⚠ 有"&{NDC}&"笔，只列前{PER_CAP}笔","")', F_RED, border=False)
    ob = f'IFERROR(INDEX({AC_OPENS},MATCH({A},{AC_NAMES},0)),0)'
    for k in range(PER_CAP):
        r = h + 1 + k
        ix = f'$X{r}'
        ws[ix] = f'={ksorted(k + 1, HD, J_R0, n, NDC)}'
        ws[ix].font = F_HELP
        g = lambda col: f'INDEX({jr(col)},{ix})'
        own = f'({g(J_ACC)}={A})'
        tr = f'({g(J_LINE)}="账户互转")'
        mir = f'AND({g(J_TO)}={A},{tr},NOT({own}))'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",{g(J_DATE)})'
        ws[f'C{r}'] = f'=IF({ix}=0,"",{g(J_MEMO)}&"")'
        ws[f'D{r}'] = f'=IF({ix}=0,"",IF(AND({own},{tr}),MAX({g(J_NET)},0),IF({mir},MAX(-{g(J_NET)},0),0)))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",IF(AND({own},{tr}),MAX(-{g(J_NET)},0),IF({mir},MAX({g(J_NET)},0),0)))'
        ws[f'F{r}'] = f'=IF({ix}=0,"",IF(AND({own},NOT({tr})),MAX(-{g(J_NET)},0),0))'
        ws[f'G{r}'] = f'=IF({ix}=0,"",IF(AND({own},NOT({tr})),MAX({g(J_NET)},0),0))'
        ws[f'H{r}'] = f'=IF({ix}=0,"",IF(AND(NOT({own}),NOT({mir}),{g(J_LINE)}="应付职工薪酬"),-{g(J_NET)},0))'
        prev = ob if k == 0 else f'I{r - 1}'
        ws[f'I{r}'] = f'=IF({ix}=0,"",ROUND(N({prev})+N(D{r})-N(E{r})-N(F{r})+N(G{r}),2))'
        ws[f'J{r}'] = f'=IF({ix}=0,"",{g(J_CT)}&"")'
        ws[f'K{r}'] = f'=IF({ix}=0,"",{g(J_PJ)}&"")'
        ws[f'L{r}'] = f'=IF({ix}=0,"",{g(J_ACC)}&"")'
        ws[f'M{r}'] = f'=IF({ix}=0,"",{g(J_NOTE)}&"")'
    r1 = h + PER_CAP
    style_rows(ws, h + 1, r1, list('ABCDEFGHIJKLM'), auto=list('ABCDEFGHIJKLM'), fmts={'B': DATE, **{c: MONEY for c in 'DEFGHI'}},
               aligns={'C': AL, 'M': AL, **{c: AR for c in 'DEFGHI'}})
    for r in range(h + 1, r1 + 1):
        for c in 'ABCDEFGHIJKLM':
            ws[f'{c}{r}'].fill = FILL_NONE
    _lbl(ws, f'H{D0 + 1}', '期初余额')
    _val(ws, f'I{D0 + 1}', f'={ob}')
    _lbl(ws, f'J{D0 + 1}', '现在余额')
    _val(ws, f'K{D0 + 1}', f'=IFERROR(INDEX({s_agg.ba(BA_E)},MATCH({A},{AC_NAMES},0)),0)', font=F_KPI_V)
    # ── 右上：保证金、押金台账（流水类别「保证金押金」，按单位） ──
    widths(ws, {'O': 5, 'P': 18, 'Q': 13, 'R': 13, 'S': 13, 'T': 11, 'U': 2, 'V': 14})
    section(ws, 3, 'O', 'V', '保证金、押金（付出去还没退回的）', C_AR)
    header(ws, 4, [('O', '序号'), ('P', '交给谁'), ('Q', '付出去'), ('R', '退回来'), ('S', '还没退'), ('T', '最近一笔'), ('V', '最近一笔的项目')], C_AR)
    nu = UN_R1 - UN_R0 + 1
    HB = 'AB'
    counter(ws, HB, 1, nu, lambda i: (f'COUNTIFS({jr(J_UN)},{SH_UNIT}!${UN_NAME}${UN_R0 + i},{jr(J_LINE)},"其他应收款",{rj})*'
                                       f'({SH_UNIT}!${UN_NAME}${UN_R0 + i}<>"")>0'))
    BN = 12
    b0 = 5
    for k in range(BN):
        r = b0 + k
        ix = f'$AC{r}'
        ws[ix] = f'={kth(k + 1, HB, 1, nu)}'
        ws[ix].font = F_HELP
        nm = f'$P{r}'
        ws[f'O{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'P{r}'] = f'=IF({ix}=0,"",INDEX({UN_NAMES},{ix}))'
        ws[f'Q{r}'] = f'=IF({ix}=0,"",-SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"其他应收款",{jr(J_NET)},"<0",{rj}))'
        ws[f'R{r}'] = f'=IF({ix}=0,"",SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"其他应收款",{jr(J_NET)},">0",{rj}))'
        ws[f'S{r}'] = f'=IF({ix}=0,"",Q{r}-R{r})'
        lp = (f'SUMPRODUCT(MAX(({jr(J_UN)}={nm})*({jr(J_LINE)}="其他应收款")*({jr(J_DN)}>0)*({jr(J_DN)}<={AX_E})'
              f'*({jr(J_DN)}*10000+ROW({jr(J_DN)}))))')
        ws[f'U{r}'] = f'=IF({ix}=0,0,{lp})'
        ws[f'U{r}'].font = F_HELP
        ws[f'T{r}'] = f'=IF(N(U{r})=0,"",INT(U{r}/10000))'
        ws[f'V{r}'] = f'=IF(N(U{r})=0,"",INDEX({SH_CASH}!${J_PJ}:${J_PJ},MOD(U{r},10000))&"")'
    r = b0 + BN
    put(ws, f'P{r}', '建账前没退的（期初）', F_TXT, align=AL)
    put(ws, f'S{r}', f'=N({oo(1)})', F_AUTO, fmt=MONEY, align=AR)
    put(ws, f'P{r + 1}', '没选「交给谁」的', F_TXT, align=AL)
    tot = f'-SUMIFS({jr(J_NET)},{jr(J_LINE)},"其他应收款",{rj})'
    put(ws, f'S{r + 1}', f'=ROUND({tot}-SUM(S{b0}:S{r - 1}),2)', F_AUTO, fmt=MONEY, align=AR)
    put(ws, f'P{r + 2}', '合计（＝资产负债表里的保证金押金）', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'S{r + 2}', f'=SUM(S{b0}:S{r + 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    for rr in (r, r + 1, r + 2):
        for c in 'OQRT':
            put(ws, f'{c}{rr}', None, fill_=FILL_TOT if rr == r + 2 else None)
    style_rows(ws, b0, b0 + BN - 1, list('OPQRSTV'), auto=list('OPQRSTV'), fmts={'Q': MONEY, 'R': MONEY, 'S': MONEY, 'T': DATE},
               aligns={'P': AL, 'Q': AR, 'R': AR, 'S': AR, 'V': AL})
    for rr in range(b0, b0 + BN):
        for c in 'OPQRSTV':
            ws[f'{c}{rr}'].fill = FILL_NONE
    put(ws, f'P{r + 3}', '投标、履约保证金、押金：流水里类别选「保证金押金」、单位选交给谁；退回来也记这个类别。', F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'P{r + 3}:T{r + 4}')
    hide(ws, 'X', 'Y', HC, HD, HB, 'AC', 'U')
    ws.freeze_panes = f'A{H0 + 1}'
    return ws


def build_all(wb, ctx):
    build_acc(wb, ctx)
    build_fund(wb, ctx)
    build_ar(wb, ctx)
    build_aps(wb, ctx)
    build_per(wb, ctx)
