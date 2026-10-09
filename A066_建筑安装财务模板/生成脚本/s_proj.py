# -*- coding: utf-8 -*-
"""项目组三张表（橙色）：
   【项目账】     选一个项目＋起止：合同收款、收入、成本、利润（本期｜开工至今），下面五张明细清单
                  （收款、应付登记、现金直接成本、考勤人工、付款），清单合计跟上面核对。
   【项目利润表】 起止：每个工程项目一行（本期｜开工至今）、合计；下面「跟公司利润表对账」（差额应为 0）。
   【费用分摊】   选年度：大白话说明、各年待摊和费率（_汇 年度表）、所选年度各项目摊多少、待摊明细。
   口径：报表口径.md §3（会计口径收入）、§7（待摊）、§8（项目成本）、§9（公司利润表，对账用）。只用定义名称取数。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *

TYPES = ('材料', '分包', '机械')
FILL_LBL = fill('FFD9E1F2')
FILL_HI = fill('FFFFF2CC')
MONTHF = 'yyyy"年"m"月"'


# ═══════════════════════════ 公式片段 ═══════════════════════════
def _dr(p, x, y, col='日期'):
    """SUMIFS/COUNTIFS 的日期段条件：p_col 在 [x, y]"""
    return f'{p}_{col},">="&{x},{p}_{col},"<="&{y}'


def _ym(d):
    return f'(YEAR({d})*100+MONTH({d}))'


def _drk(x, y):
    """考勤按月：考_年月 在 [x 所在月, y 所在月]"""
    return f'考_年月,">="&{_ym(x)},考_年月,"<="&{_ym(y)}'


def _pay(pe, x, y, typ=None):
    """应付登记（按应付日期）"""
    t = f',付_类型,"{typ}"' if typ else ''
    return f'SUMIFS(付_应付额,付_项目,{pe}{t},{_dr("付", x, y)})'


def _cash(pe, x, y, cls):
    """收支登记里归属＝项目成本的现金成本（正数）；前面自带负号"""
    return f'-SUMIFS(收_净额,收_项目,{pe},收_成本类,"{cls}",收_归属,"项目成本",{_dr("收", x, y)})'


def _kq(pe, x, y, col='金额', jt=True):
    """考勤 8 对按项目：Σi SUMIFS(考_{col}i, 考_项目i, 项目)"""
    c = ',考_计提,1' if jt else ''
    return '(' + '+'.join(f'SUMIFS(考_{col}{i + 1},考_项目{i + 1},{pe}{c},{_drk(x, y)})' for i in range(NPAIR)) + ')'


def costs(pe, x, y):
    """项目（pe＝已转义的项目名）在 [x, y] 的直接成本、税金、分摊（报表口径 §8）。
       其他直接费＝应付登记里不是材料/分包/机械的（「其他」或写错的类型）＋现金其他直接费＋归属项目成本的管理费用类开支"""
    m = {t: _pay(pe, x, y, t) + _cash(pe, x, y, t) for t in TYPES}
    m['人工'] = _kq(pe, x, y) + _cash(pe, x, y, '人工')
    m['其他直接费'] = (_pay(pe, x, y) + ''.join(f'-{_pay(pe, x, y, t)}' for t in TYPES)
                    + _cash(pe, x, y, '其他直接费') + _cash(pe, x, y, '管理费用'))
    m['税金'] = (f'(SUMIFS(应_销项税,应_项目,{pe},{_dr("应", x, y)})'
               f'-SUMIFS(付_进项税,付_项目,{pe},{_dr("付", x, y, "开票日期")}))*(1+P_附加税率)')
    m['分摊'] = (f'SUMIFS(收_摊费,收_项目,{pe},{_dr("收", x, y)})+SUMIFS(付_摊费,付_项目,{pe},{_dr("付", x, y)})+'
               + _kq(pe, x, y, '摊', jt=False))
    return m


def rev_cum(pe, d):
    """项目到 d 为止的累计确认收入（会计口径）"""
    return rev_formula(cum_kp(pe, d), cum_sk(pe, d), cum_cz(pe, d))


def _hid(ws, cell, value):
    ws[cell] = value
    ws[cell].font = F_HELP


def _lbl(ws, cell, text):
    put(ws, cell, text, F_KPI_L, FILL_LBL, align=AC)


def _list_area(ws, r0, n, cols, fmts=None, aligns=None, key=None):
    """清单显示区：没数的行不画框（条件格式：有数才画框）"""
    fmts, aligns = fmts or {}, aligns or {}
    for r in range(r0, r0 + n):
        for c in cols:
            cell = ws[f'{c}{r}']
            cell.font = F_AUTO
            cell.alignment = aligns.get(c, AC)
            if c in fmts:
                cell.number_format = fmts[c]
    key = key or cols[0]
    ws.conditional_formatting.add(f'{cols[0]}{r0}:{cols[-1]}{r0 + n - 1}',
                                  FormulaRule(formula=[f'${key}{r0}<>""'], border=BD))


def _note_rows(ws, r0, c1, c2, notes, height=30):
    for i, t in enumerate(notes):
        r = r0 + i
        ws.merge_cells(f'{c1}{r}:{c2}{r}')
        put(ws, f'{c1}{r}', t, F_TXT, align=ALW, border=False)
        ws.row_dimensions[r].height = height


# ═══════════════════════════ 项目账 ═══════════════════════════
BK_NR = 300                       # 每张清单最多列几条
BK_L0 = 32                        # 清单第一行
BK = dict(P='$AJ$3', PE='$AJ$4', IX='$AJ$5', S='$AJ$6', X='$AJ$7', Y='$AJ$8', TP='$AJ$9')
BK_ROWS = ['合同额', '应收总额（合同口径）', '已开票', '已收', '未收', '确认收入（会计口径）', '材料', '分包', '机械', '人工',
           '其他直接费', '直接成本合计', '税金（估算）', '毛利', '分摊管理费', '净利', '净利率']
BK_R = {k: 9 + i for i, k in enumerate(BK_ROWS)}          # ① 汇总：第 9～25 行；C＝本期，D＝开工至今


def build_book(wb, ctx):
    ws = wb[SH_PJBOOK]
    W = {'A': 11, 'B': 12, 'C': 15, 'D': 15, 'E': 30, 'F': 2, 'G': 11, 'H': 14, 'I': 7, 'J': 22, 'K': 13, 'L': 13, 'M': 2,
         'N': 11, 'O': 12, 'P': 13, 'Q': 24, 'R': 13, 'S': 11, 'T': 2, 'U': 10, 'V': 9, 'W': 7, 'X': 13, 'Y': 24, 'Z': 2,
         'AA': 11, 'AB': 12, 'AC': 13, 'AD': 14, 'AE': 22, 'AF': 13, 'AG': 14}
    widths(ws, W)
    tip = ('💡 三个黄格可以改：项目（下拉选）、起、止。起空着＝这个项目的开工日（【项目档案】开工日期，没填＝建账日），止空着＝首页截止日；'
           '「开工至今」总是从开工日算到止。材料、分包、机械以【应付登记】为准（送一车料、结一次账就记一笔，不管付没付钱），'
           '付给这些供应商的钱是在还欠款，只列在最右边「付款」里看，不再算成本；人工按【考勤工资】算。下面五张清单按日期排，合计跟上面对得上（看「核对」那行）。')
    title(ws, '项目账（选一个项目：收入、成本、利润和每一笔明细）', 'L', C_PROJ, tip=tip)
    P, PE, IX, S, X, Y, TP = (BK[k] for k in ('P', 'PE', 'IX', 'S', 'X', 'Y', 'TP'))

    # ── 隐藏的「实际用的值」（AI 说明、AJ 值）
    pre = [f'COUNTIFS(收_项目,{PE},收_有效,1,收_日期,"<"&{S})',
           f'COUNTIFS(付_项目,{PE},付_有效,1,付_日期,">0",付_日期,"<"&{S})',
           f'COUNTIFS(应_项目,{PE},应_有效,1,应_日期,">0",应_日期,"<"&{S})',
           '+'.join(f'COUNTIFS(考_项目{i + 1},{PE},考_有效,1,考_天数{i + 1},"<>0",考_年月,"<"&{_ym(S)})' for i in range(NPAIR))]
    hid = [('项目', f'=IF(TRIM($B$3&"")<>"",TRIM($B$3&""),IFERROR(INDEX(项目_名称,MATCH("工程",项目_类型,0))&"",""))'),
           ('转义', f'={esc(P)}'),
           ('第几个', f'=IF({P}="",0,IFERROR(MATCH({P},项目_名称,0),0))'),
           ('开工有效', f'=IF({IX}=0,P_建账日,INDEX(项目_开工有效,{IX}))'),
           ('起', f'=IF(ISNUMBER($B$4),INT($B$4),{S})'),
           ('止', f'=IF(ISNUMBER($B$5),INT($B$5),P_截止)'),
           ('类型', f'=IF({IX}=0,"",INDEX(项目_类型,{IX})&"")'),
           ('开工前收支', f'=IF({P}="",0,{pre[0]})'), ('开工前应付', f'=IF({P}="",0,{pre[1]})'),
           ('开工前应收', f'=IF({P}="",0,{pre[2]})'), ('开工前考勤', f'=IF({P}="",0,{pre[3]})'),
           ('开工前合计', '=SUM(AJ10:AJ13)')]
    for i, (k, f) in enumerate(hid):
        _hid(ws, f'AI{3 + i}', k)
        _hid(ws, f'AJ{3 + i}', f)

    # ── 选择格
    selector(ws, 'A3', '项目', 'B3', None, '=项目列表', prompt='选项目；空着＝第一个工程项目')
    selector(ws, 'A4', '起', 'B4', None, fmt=DATE)
    selector(ws, 'A5', '止', 'B5', None, fmt=DATE)
    for r in (3, 4, 5):
        ws.merge_cells(f'B{r}:C{r}')
        ws.merge_cells(f'D{r}:E{r}')
    dv_date(ws, 'B4')
    dv_date(ws, 'B5')
    put(ws, 'D3', f'=IF({P}="","⚠ 【项目档案】里还没有工程项目","实际看："&{P}&IF(TRIM($B$3&"")="","（空＝第一个工程项目）",""))',
        F_NOTE, align=AL, border=False)
    put(ws, 'D4', f'="实际："&TEXT({X},"yyyy/mm/dd")&IF(ISNUMBER($B$4),"","（空＝开工日）")', F_NOTE, align=AL, border=False)
    put(ws, 'D5', f'="实际："&TEXT({Y},"yyyy/mm/dd")&IF(ISNUMBER($B$5),"","（空＝首页截止日）")', F_NOTE, align=AL, border=False)
    # 项目信息
    _lbl(ws, 'G3', '甲方')
    ws.merge_cells('H3:J3')
    put(ws, 'H3', f'=IF({IX}=0,"",INDEX(项目_甲方,{IX})&"")', F_AUTOB, FILL_AUTO, align=AL)
    _lbl(ws, 'K3', '合同额')
    put(ws, 'L3', f'=IF({IX}=0,0,INDEX(项目_合同,{IX}))', F_AUTOB, FILL_AUTO, MONEY, AR)
    _lbl(ws, 'G4', '开工')
    put(ws, 'H4', f'=IF({IX}=0,"",IF(INDEX(项目_开工,{IX})>0,INDEX(项目_开工,{IX}),"没填（按建账日）"))', F_AUTOB, FILL_AUTO, DATE, AC)
    _lbl(ws, 'I4', '完工')
    put(ws, 'J4', f'=IF({IX}=0,"",IF(INDEX(项目_完工,{IX})>0,INDEX(项目_完工,{IX}),"还没完工"))', F_AUTOB, FILL_AUTO, DATE, AC)
    _lbl(ws, 'K4', '状态')
    put(ws, 'L4', f'=IF({IX}=0,"",INDEX(项目_状态,{IX})&"")', F_AUTOB, FILL_AUTO, align=AC)
    ws.merge_cells('G5:L5')
    put(ws, 'G5', f'=IF({P}="","",IF({IX}=0,"⚠ 「"&{P}&"」不在【项目档案】里",IF({TP}<>"工程","⚠ 「"&{P}&"」不是工程项目（类型＝"&{TP}'
                  f'&"），它的开支进待摊，请看【费用分摊】",IF({X}>{Y},"⚠ 起 晚于 止，请改日期",""))))', F_RED, align=AL, border=False)
    ws.merge_cells('A6:L6')
    put(ws, 'A6', (f'=IF($AJ$14=0,"","⚠ 开工日 "&TEXT({S},"yyyy/mm/dd")&" 以前还有 "&$AJ$14&" 处记在这个项目上（收支 "&$AJ$10&" 笔、应付 "&$AJ$11'
                   f'&" 笔、应收 "&$AJ$12&" 笔、考勤 "&$AJ$13&" 处），「开工至今」不含它们；如果开工日填晚了，请到【项目档案】改开工日期。")'),
        F_RED, align=AL, border=False)
    home_link(ws, 'N3')

    # ── ① 汇总
    section(ws, 7, 'A', 'E', '① 这个项目赚了多少', C_PROJ)
    header(ws, 8, [('A', '项目'), ('C', f'="本期 "&TEXT({X},"yyyy/m/d")&"～"&TEXT({Y},"yyyy/m/d")'),
                   ('D', f'="开工至今 "&TEXT({S},"yyyy/m/d")&"～"&TEXT({Y},"yyyy/m/d")'), ('E', '怎么算')], C_PROJ, height=40)
    ws.merge_cells('A8:B8')
    R = BK_R
    g = lambda f: f'=IF({P}="",0,{f})'
    col_f = {}
    for col, x in (('C', X), ('D', S)):
        m = costs(PE, x, Y)
        f = {'已开票': g(f'SUMIFS(应_开票额,应_项目,{PE},{_dr("应", x, Y)})'),
             '已收': g(f'SUMIFS(收_净额,收_归类,"工程款收款",收_项目,{PE},{_dr("收", x, Y)})'),
             '确认收入（会计口径）': g(f'({rev_cum(PE, Y)})-({rev_cum(PE, f"({x}-1)")})'),
             '材料': g(m['材料']), '分包': g(m['分包']), '机械': g(m['机械']), '人工': g(m['人工']),
             '其他直接费': g(m['其他直接费']),
             '直接成本合计': f'=SUM({col}{R["材料"]}:{col}{R["其他直接费"]})',
             '税金（估算）': g(m['税金']),
             '毛利': f'={col}{R["确认收入（会计口径）"]}-{col}{R["直接成本合计"]}-{col}{R["税金（估算）"]}',
             '分摊管理费': g(m['分摊']),
             '净利': f'={col}{R["毛利"]}-{col}{R["分摊管理费"]}',
             '净利率': f'=IF(N({col}{R["确认收入（会计口径）"]})=0,"",{col}{R["净利"]}/{col}{R["确认收入（会计口径）"]})'}
        col_f[col] = f
    col_f['D']['合同额'] = f'=IF({IX}=0,0,INDEX(项目_合同,{IX}))'
    col_f['D']['应收总额（合同口径）'] = g(f'D{R["合同额"]}+SUMIFS(应_应收额,应_项目,{PE},应_日期,"<="&{Y})'
                                     f'+SUMIFS(期初_金额,期初_类型,"应收账款",期初_项目,{PE})')
    col_f['D']['未收'] = f'=D{R["应收总额（合同口径）"]}-D{R["已收"]}'
    notes = {'合同额': '【项目档案】合同金额', '应收总额（合同口径）': '合同＋补充、签证、调整−扣款＋期初应收',
             '已开票': '【应收登记】开票', '已收': '工程款（含专户、个人账户收的）', '未收': '应收总额−已收',
             '确认收入（会计口径）': '=IF(P_收入口径="开票与收款取大","有确认产值按产值，否则开票、收款取大","按收入口径「"&P_收入口径&"」")',
             '材料': '应付登记的材料＋直接付钱买的', '分包': '应付登记的分包＋直接付的',
             '机械': '应付登记的台班＋直接付的',
             '人工': '="考勤（"&TEXT(P_考勤起算,"yyyy年m月")&"起）＋没考勤的人发的工资"',
             '其他直接费': '项目其他费用＋应付登记「其他」', '直接成本合计': '上面五项',
             '税金（估算）': '销项税−专票进项税，再加附加税', '毛利': '确认收入−直接成本−税金',
             '分摊管理费': '施工费×当年费率，见【费用分摊】', '净利': '毛利−分摊管理费', '净利率': '净利÷确认收入'}
    bal = ('合同额', '应收总额（合同口径）', '未收')
    strong = ('确认收入（会计口径）', '直接成本合计', '毛利', '净利')
    for k, r in R.items():
        b = k in strong
        fl = FILL_TOT if k == '净利' else (FILL_SUB if b else None)
        ws.merge_cells(f'A{r}:B{r}')
        put(ws, f'A{r}', k, F_TXTB if b else F_TXT, fl, align=AL)
        ws[f'B{r}'].border = BD
        fmt = PCT if k == '净利率' else MONEY
        for col in ('C', 'D'):
            if col == 'C' and k in bal:
                put(ws, f'C{r}', '—', F_NOTE, FILL_AUTO, align=AC)
                continue
            put(ws, f'{col}{r}', col_f[col][k], F_AUTOB if b else F_AUTO, fl, fmt, AC if k == '净利率' else AR)
        put(ws, f'E{r}', notes[k], F_NOTE, align=AL)
    ws.conditional_formatting.add(f'C{R["净利"]}:D{R["净利"]}', FormulaRule(formula=[f'N(C{R["净利"]})<0'], fill=FILL_WARN))
    # 右边：怎么算的（大白话）
    section(ws, 7, 'G', 'L', '这些数怎么来的', C_PROJ)
    expl = ['「开工至今」从【项目档案】的开工日算到止日（开工日没填＝建账日）；「本期」按上面黄格的起止。改了黄格马上就变。',
            '收入：甲方确认了产值的按产值；没有的，开票和收款哪个多按哪个（【基本信息】可以改口径）。合同、应收、未收是到止日的余额，只看右边一列。',
            '材料、分包、机械：以【应付登记】为准（送货单、结算单、台班记一笔算一笔，不管付没付钱）。付给这些供应商的钱是还欠款，列在「付款」清单，不再算成本；'
            '没登记应付、直接付钱买的零星材料，按【收支登记】算。',
            '=TEXT(P_考勤起算,"yyyy年m月")&"起，人工按【考勤工资】算（这个项目的天数×单价，月薪按天数比例分），发工资只是还欠薪；在那以前、或者没录考勤的人，按实际发的工资算。"',
            '其他直接费：【收支登记】里记在这个项目上的项目其他费用（运费、检测、零星开支等）和【应付登记】类型为「其他」的。',
            '税金是估算：开给甲方发票的销项税 − 收到的专票进项税，再加附加税（【基本信息】的税率）。',
            '分摊管理费：办公室开支、管理人员在公司的工资等「待摊费用」，每年按各项目的施工费（人工＋分包＋机械）分摊，见【费用分摊】。',
            '下面五张清单只列「本期」（起～止）的，按日期排；每张最多列 300 条，合计跟上面对得上。']
    for i, t in enumerate(expl):
        r = 8 + 2 * i
        ws.merge_cells(f'G{r}:L{r + 1}')
        put(ws, f'G{r}', t, F_TXT, align=ALW, border=False)
    # 核对
    ws.merge_cells('A26:E26')
    put(ws, 'A26', (f'=IF({P}="","",IF(AND(ROUND($D$30-$C${R["已收"]},2)=0,ROUND($K$30+$R$30+$X$30-$C${R["直接成本合计"]},2)=0),'
                    f'"✓ 核对：下面清单合计跟「本期」对得上（收款＝已收；应付登记＋现金直接成本＋考勤人工＝直接成本合计）",'
                    f'"✗ 核对：下面清单合计跟「本期」对不上（每张清单最多列 {BK_NR} 条，超出请缩短起止日期）"))'),
        F_AUTOB, align=AL, border=False)

    # ── ② 明细清单（本期）
    L0, NR = BK_L0, BK_NR
    L1 = L0 + NR - 1
    section(ws, 28, 'A', 'AG', '② 明细（本期：起～止，按日期排；只看，要改请到录入表）', C_PROJ)
    # 隐藏：BA 第几条（常数）、BB 收支三张清单共用排序键（类型×10^9＋日期×10000＋第几条）、BC 应付排序键、BD 考勤排序键
    for i in range(N_CASH):
        _hid(ws, f'BA{2 + i}', i + 1)
    for i in range(N_CASH):
        r = 2 + i
        n = f'$BA{r}'
        ws[f'BB{r}'] = (f'=IF(OR({P}="",INDEX(收_有效,{n})<>1,INDEX(收_项目,{n})<>{P},INDEX(收_日期,{n})<{X},INDEX(收_日期,{n})>{Y}),"",'
                        f'IF(INDEX(收_归类,{n})="工程款收款",1,IF(INDEX(收_归属,{n})="项目成本",2,'
                        f'IF(INDEX(收_冲应付,{n})+INDEX(收_冲工资,{n})>0,3,0)))*1000000000+INDEX(收_日期,{n})*10000+{n})')
        ws[f'BB{r}'].font = F_HELP
    for i in range(N_AP):
        r = 2 + i
        n = f'$BA{r}'
        ws[f'BC{r}'] = (f'=IF(OR({P}="",INDEX(付_有效,{n})<>1,INDEX(付_项目,{n})<>{P},INDEX(付_日期,{n})<{X},INDEX(付_日期,{n})>{Y}),"",'
                        f'INDEX(付_日期,{n})*10000+{n})')
        ws[f'BC{r}'].font = F_HELP
    for i in range(N_ATT):
        r = 2 + i
        n = f'$BA{r}'
        days = '+'.join(f'(INDEX(考_项目{j + 1},{n})={P})*INDEX(考_天数{j + 1},{n})' for j in range(NPAIR))
        ws[f'BD{r}'] = (f'=IF(OR({P}="",INDEX(考_有效,{n})<>1,INDEX(考_年月,{n})<{_ym(X)},INDEX(考_年月,{n})>{_ym(Y)}),"",'
                        f'IF({days}<>0,INDEX(考_月,{n})*10000+{n},""))')
        ws[f'BD{r}'].font = F_HELP
    KB, KC, KD = f'$BB$2:$BB${N_CASH + 1}', f'$BC$2:$BC${N_AP + 1}', f'$BD$2:$BD${N_ATT + 1}'
    cnts = [('收款前', f'=COUNTIF({KB},"<1000000000")'), ('收款', f'=COUNTIFS({KB},">=1000000000",{KB},"<2000000000")'),
            ('现金前', f'=COUNTIF({KB},"<2000000000")'), ('现金', f'=COUNTIFS({KB},">=2000000000",{KB},"<3000000000")'),
            ('付款前', f'=COUNTIF({KB},"<3000000000")'), ('付款', f'=COUNTIFS({KB},">=3000000000",{KB},"<4000000000")'),
            ('应付', f'=COUNT({KC})'), ('考勤', f'=COUNT({KD})')]
    CN = {}
    for i, (k, f) in enumerate(cnts):
        _hid(ws, f'BE{3 + i}', k)
        _hid(ws, f'BF{3 + i}', f)
        CN[k] = f'$BF${3 + i}'
    for i in range(NR):
        r = L0 + i
        k = f'$BM{r}'
        _hid(ws, f'BM{r}', i + 1)
        for col, nm, src in (('BH', '收款', KB), ('BJ', '现金', KB), ('BL', '付款', KB)):
            _hid(ws, f'{col}{r}', f'=IF({k}>{CN[nm]},0,MOD(SMALL({src},{CN[nm + "前"]}+{k}),10000))')
        _hid(ws, f'BI{r}', f'=IF({k}>{CN["应付"]},0,MOD(SMALL({KC},{k}),10000))')
        _hid(ws, f'BK{r}', f'=IF({k}>{CN["考勤"]},0,MOD(SMALL({KD},{k}),10000))')

    def cell(r, ix, expr, text=False):
        return f'=IF({ix}=0,"",{expr}{"&" + chr(34) * 2 if text else ""})'

    lists = [
        ('收款（工程款，含专户、个人账户收的）', 'A', 'E', '$BH', CN['收款'], C_AR,
         [('A', '日期', '收_日期', DATE, 0), ('B', '账户', '收_账户', None, 1), ('C', '客户', '收_客户', None, 1),
          ('D', '金额', '收_净额', MONEY, 0), ('E', '摘要', '收_摘要', None, 1)], ['D']),
        ('应付登记（材料、分包、机械、其他：记了就算成本）', 'G', 'L', '$BI', CN['应付'], C_AR,
         [('G', '日期', '付_日期', DATE, 0), ('H', '供应商', '付_供应商', None, 1), ('I', '类型', '付_类型', None, 1),
          ('J', '摘要', '付_摘要', None, 1), ('K', '应付金额', '付_应付额', MONEY, 0), ('L', '已开票', '付_已开票', MONEY, 0)], ['K', 'L']),
        ('现金直接成本（没登记应付、直接付钱的）', 'N', 'S', '$BJ', CN['现金'], C_CASH,
         [('N', '日期', '收_日期', DATE, 0), ('O', '账户', '收_账户', None, 1), ('P', '收支项目', '收_收支项目', None, 1),
          ('Q', '摘要', '收_摘要', None, 1), ('R', '金额', None, MONEY, 0), ('S', '算进', None, None, 0)], ['R']),
        ('考勤人工（这个项目的天数和工资）', 'U', 'Y', '$BK', CN['考勤'], C_HOME,
         [('U', '月份', '考_月', MONTHF, 0), ('V', '姓名', '考_姓名', None, 1), ('W', '天数', None, '0.##', 0),
          ('X', '进成本的工资', None, MONEY, 0), ('Y', '说明', None, None, 0)], ['W', 'X']),
        ('付款（还应付款、发工资：只看，不算成本）', 'AA', 'AG', '$BL', CN['付款'], C_VIEW,
         [('AA', '日期', '收_日期', DATE, 0), ('AB', '账户', '收_账户', None, 1), ('AC', '收支项目', '收_收支项目', None, 1),
          ('AD', '付给', None, None, 0), ('AE', '摘要', '收_摘要', None, 1), ('AF', '金额', None, MONEY, 0), ('AG', '算作', None, None, 0)], ['AF']),
    ]
    pj_days = lambda ix, what: '+'.join(f'(INDEX(考_项目{j + 1},{ix})={P})*INDEX(考_{what}{j + 1},{ix})' for j in range(NPAIR))
    for ttl, c1, c2, ixc, ncell, color, cols, sums in lists:
        section(ws, 29, c1, c2, ttl, color)
        header(ws, 31, [(c, t) for c, t, *_ in cols], color, height=30)
        first_sum = sums[0]
        cnt_end = CL(CI(first_sum) - 1)
        ws.merge_cells(f'{c1}30:{cnt_end}30')
        put(ws, f'{c1}30', f'="共 "&{ncell}&" 条"&IF({ncell}>{NR},"，只列前 {NR} 条（请缩短起止日期）","")&"　合计 →"',
            F_NOTE, align=AR)
        for c in sums:
            put(ws, f'{c}30', f'=SUM({c}{L0}:{c}{L1})', F_AUTOB, FILL_TOT, MONEY if c != 'W' else '0.##', AR)
        for c in [c for c, *_ in cols if c not in sums and CI(c) > CI(cnt_end)]:
            put(ws, f'{c}30', None, fill_=FILL_TOT)
        for i in range(NR):
            r = L0 + i
            ix = f'{ixc}{r}'
            for c, t, nm, fmt, txt in cols:
                if nm:
                    ws[f'{c}{r}'] = cell(r, ix, f'INDEX({nm},{ix})', txt)
            if c1 == 'N':
                ws[f'R{r}'] = cell(r, ix, f'-INDEX(收_净额,{ix})')
                ws[f'S{r}'] = cell(r, ix, f'IF(INDEX(收_成本类,{ix})="管理费用","其他直接费",INDEX(收_成本类,{ix})&"")')
            elif c1 == 'U':
                ws[f'W{r}'] = cell(r, ix, pj_days(ix, '天数'))
                ws[f'X{r}'] = cell(r, ix, f'IF(INDEX(考_计提,{ix})=1,{pj_days(ix, "金额")},0)')
                ws[f'Y{r}'] = cell(r, ix, f'IF(INDEX(考_计提,{ix})=1,"",IF(INDEX(考_过账,{ix})=1,"过账人员，不进成本",'
                                          f'"不进成本（考勤起算月以前，或人工按实际发放）"))')
            elif c1 == 'AA':
                ws[f'AD{r}'] = cell(r, ix, f'IF(INDEX(收_冲应付,{ix})=1,INDEX(收_供应商,{ix}),INDEX(收_人员,{ix}))&""')
                ws[f'AF{r}'] = cell(r, ix, f'-INDEX(收_净额,{ix})')
                ws[f'AG{r}'] = cell(r, ix, f'IF(INDEX(收_冲应付,{ix})=1,"还应付款","发工资（还欠薪）")')
        _list_area(ws, L0, NR, [c for c, *_ in cols], {c: f for c, _t, _n, f, _x in cols if f},
                   {**{c: AL for c, t, *_ in cols if t in ('摘要', '供应商', '付给', '说明', '收支项目', '客户', '账户')},
                    **{c: AR for c, _t, _n, f, _x in cols if f in (MONEY,)}})
    hide(ws, 'AI', 'AJ', 'BA', 'BB', 'BC', 'BD', 'BE', 'BF', 'BH', 'BI', 'BJ', 'BK', 'BL', 'BM')
    ws.freeze_panes = 'A6'
    print_setup(ws, '1:5', landscape=True)
    return ws


# ═══════════════════════════ 项目利润表 ═══════════════════════════
PL_NP, PL_R0 = 60, 8               # 最多列 60 个工程项目；第 8 行起；第 7 行合计
PL_BLK = ['收入', '材料', '分包', '机械', '人工', '其他直接费', '税金', '毛利', '分摊管理费', '净利', '净利率']
PL_C1 = {k: CL(7 + i) for i, k in enumerate(PL_BLK)}        # 本期：G～Q
PL_C2 = {k: CL(18 + i) for i, k in enumerate(PL_BLK)}       # 开工至今：R～AB
PL_TIP = 'AC'
PL_X, PL_Y = '$AW$3', '$AW$4'
PL_RB = PL_R0 + PL_NP + 2                                     # 对账块的标题行


def build_pl(wb, ctx):
    ws = wb[SH_PJPL]
    NP, R0 = PL_NP, PL_R0
    R1 = R0 + NP - 1
    W = {'A': 5, 'B': 15, 'C': 7, 'D': 11, 'E': 11, 'F': 13, PL_TIP: 42}
    for k in PL_BLK:
        W[PL_C1[k]] = W[PL_C2[k]] = 8 if k == '净利率' else 13
    widths(ws, W)
    tip = ('💡 黄格是本期的起、止（空着＝首页的年初～截止日）。每个工程项目一行：左边「本期」，右边「开工至今」（各项目从自己的开工日算到止日）。'
           '收入按会计口径（确认产值，或开票、收款取大）；材料、分包、机械按【应付登记】＋没登记应付直接付的；人工按考勤；分摊管理费见【费用分摊】。'
           '最下面跟公司利润表对账，差额应为 0。')
    title(ws, '项目利润表（每个工程项目赚了多少：本期＋开工至今）', PL_C1['净利率'], C_PROJ, tip=tip)
    X, Y = PL_X, PL_Y
    _hid(ws, 'AV3', '起')
    _hid(ws, 'AW3', '=IF(ISNUMBER($C$3),INT($C$3),P_年初)')
    _hid(ws, 'AV4', '止')
    _hid(ws, 'AW4', '=IF(ISNUMBER($F$3),INT($F$3),P_截止)')
    selector(ws, 'B3', '起', 'C3', None, fmt=DATE)
    ws.merge_cells('C3:D3')
    selector(ws, 'E3', '止', 'F3', None, fmt=DATE)
    dv_date(ws, 'C3')
    dv_date(ws, 'F3')
    ws.merge_cells('G3:Q3')
    put(ws, 'G3', (f'="实际："&TEXT({X},"yyyy/mm/dd")&"～"&TEXT({Y},"yyyy/mm/dd")&"（空着＝首页的年初～截止日）"'
                   f'&IF({X}>{Y},"　⚠ 起 晚于 止，请改日期","")'), F_NOTE, align=AL, border=False)
    home_link(ws, f'{PL_TIP}3')
    # 只列工程项目
    counter(ws, 'AU', R0, N_PJ, lambda i: f'INDEX(项目_类型,{i + 1})="工程"')
    NCNT = cnt('AU', R0, N_PJ)
    ws.merge_cells('A4:F4')
    put(ws, 'A4', f'="共 "&{NCNT}&" 个工程项目"&IF({NCNT}>{NP},"，只列前 {NP} 个（合计和对账会对不上）","")', F_NOTE, align=AL, border=False)
    ws.merge_cells('G4:L4')
    c = put(ws, 'G4', f'↓ 跟公司利润表对账（第 {PL_RB} 行）', Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'),
            align=AL, border=False)
    link(c, SH_PJPL, f'A{PL_RB}')
    # 表头
    ws.merge_cells('A5:F5')
    put(ws, 'A5', '项目', F_HDR, fill(C_PROJ), align=AC)
    c1a, c1b = PL_C1['收入'], PL_C1['净利率']
    c2a, c2b = PL_C2['收入'], PL_C2['净利率']
    ws.merge_cells(f'{c1a}5:{c1b}5')
    put(ws, f'{c1a}5', f'="本期 "&TEXT({X},"yyyy/m/d")&"～"&TEXT({Y},"yyyy/m/d")', F_HDR, fill(C_PROJ), align=AC)
    ws.merge_cells(f'{c2a}5:{c2b}5')
    put(ws, f'{c2a}5', f'="开工至今（各项目开工日～"&TEXT({Y},"yyyy/m/d")&"）"', F_HDR, fill('FF843C0C'), align=AC)
    put(ws, f'{PL_TIP}5', '', F_HDR, fill(C_PROJ), align=AC)
    heads = [('A', '序号'), ('B', '项目'), ('C', '状态'), ('D', '开工'), ('E', '完工'), ('F', '合同额')]
    heads += [(PL_C1[k], k) for k in PL_BLK] + [(PL_C2[k], k) for k in PL_BLK] + [(PL_TIP, '提示')]
    header(ws, 6, heads, C_PROJ, height=30)
    for k in PL_BLK:
        ws[f'{PL_C2[k]}6'].fill = fill('FF843C0C')
    # 每个项目一行
    # 隐藏：AE k、AF 第几个、AG 转义名、AH 开工有效、AI～AQ 累计开票/收款/产值（止、起前、开工前）、AR 本期应付登记、AS 本期进项税、AT 开工前笔数
    cum = [('AI', 'kp', Y), ('AJ', 'sk', Y), ('AK', 'cz', Y), ('AL', 'kp', f'({X}-1)'), ('AM', 'sk', f'({X}-1)'), ('AN', 'cz', f'({X}-1)'),
           ('AO', 'kp', '($AH{r}-1)'), ('AP', 'sk', '($AH{r}-1)'), ('AQ', 'cz', '($AH{r}-1)')]
    fn = {'kp': cum_kp, 'sk': cum_sk, 'cz': cum_cz}
    for i in range(NP):
        r = R0 + i
        ix, nm, pe, s = f'$AF{r}', f'$B{r}', f'$AG{r}', f'$AH{r}'
        _hid(ws, f'AE{r}', i + 1)
        _hid(ws, f'AF{r}', '=' + kth(f'$AE{r}', 'AU', R0, N_PJ))
        _hid(ws, f'AG{r}', f'=IF({nm}="","",{esc(nm)})')
        _hid(ws, f'AH{r}', f'=IF({ix}=0,0,INDEX(项目_开工有效,{ix}))')
        for col, k, d in cum:
            _hid(ws, f'{col}{r}', f'=IF({nm}="",0,{fn[k](pe, d.format(r=r))})')
        _hid(ws, f'AR{r}', f'=IF({nm}="",0,{_pay(pe, X, Y)})')
        _hid(ws, f'AS{r}', f'=IF({nm}="",0,SUMIFS(付_进项税,付_项目,{pe},{_dr("付", X, Y, "开票日期")}))')
        pre = (f'COUNTIFS(收_项目,{pe},收_有效,1,收_日期,"<"&{s})+COUNTIFS(付_项目,{pe},付_有效,1,付_日期,">0",付_日期,"<"&{s})'
               f'+COUNTIFS(应_项目,{pe},应_有效,1,应_日期,">0",应_日期,"<"&{s})+'
               + '+'.join(f'COUNTIFS(考_项目{j + 1},{pe},考_有效,1,考_天数{j + 1},"<>0",考_年月,"<"&{_ym(s)})' for j in range(NPAIR)))
        _hid(ws, f'AT{r}', f'=IF({nm}="",0,{pre})')
        put(ws, f'A{r}', f'=IF({ix}=0,"",$AE{r})', F_AUTO, align=AC, border=False)
        put(ws, f'B{r}', f'=IF({ix}=0,"",INDEX(项目_名称,{ix})&"")', F_TXTB, align=AL, border=False)
        put(ws, f'C{r}', f'=IF({ix}=0,"",INDEX(项目_状态,{ix})&"")', F_AUTO, align=AC, border=False)
        put(ws, f'D{r}', f'=IF({ix}=0,"",{s})', F_AUTO, fmt=DATE, align=AC, border=False)
        put(ws, f'E{r}', f'=IF({ix}=0,"",IF(INDEX(项目_完工,{ix})>0,INDEX(项目_完工,{ix}),""))', F_AUTO, fmt=DATE, align=AC, border=False)
        put(ws, f'F{r}', f'=IF({ix}=0,"",INDEX(项目_合同,{ix}))', F_AUTO, fmt=MONEY, align=AR, border=False)
        for C, x, kc in ((PL_C1, X, ('AL', 'AM', 'AN')), (PL_C2, s, ('AO', 'AP', 'AQ'))):
            m = costs(pe, x, Y)
            g = lambda f: f'=IF({nm}="","",{f})'
            rb = rev_formula(f'$AI{r}', f'$AJ{r}', f'$AK{r}')
            ra = rev_formula(*(f'${c}{r}' for c in kc))
            fs = {'收入': g(f'({rb})-({ra})'), '材料': g(m['材料']), '分包': g(m['分包']), '机械': g(m['机械']), '人工': g(m['人工']),
                  '其他直接费': g(m['其他直接费']), '税金': g(m['税金']),
                  '毛利': g(f'{C["收入"]}{r}-SUM({C["材料"]}{r}:{C["其他直接费"]}{r})-{C["税金"]}{r}'),
                  '分摊管理费': g(m['分摊']), '净利': g(f'{C["毛利"]}{r}-{C["分摊管理费"]}{r}'),
                  '净利率': f'=IF(OR({nm}="",N({C["收入"]}{r})=0),"",{C["净利"]}{r}/{C["收入"]}{r})'}
            for k in PL_BLK:
                b = k in ('毛利', '净利')
                put(ws, f'{C[k]}{r}', fs[k], F_AUTOB if b else F_AUTO, fmt=PCT if k == '净利率' else MONEY,
                    align=AC if k == '净利率' else AR, border=False)
        put(ws, f'{PL_TIP}{r}', f'=IF(OR({nm}="",$AT{r}=0),"","开工日前还有 "&$AT{r}&" 处记在本项目（开工至今不含）")',
            F_NOTE, align=AL, border=False)
    last = PL_TIP
    ws.conditional_formatting.add(f'A{R0}:{last}{R1}', FormulaRule(formula=[f'$B{R0}<>""'], border=BD))
    for k in ('净利', '毛利'):
        for C in (PL_C1, PL_C2):
            ws.conditional_formatting.add(f'{C[k]}{R0}:{C[k]}{R1}', FormulaRule(formula=[f'AND(ISNUMBER({C[k]}{R0}),{C[k]}{R0}<0)'], fill=FILL_WARN))
    # 合计行（第 7 行）
    rt = R0 - 1
    ws.merge_cells(f'A{rt}:E{rt}')
    put(ws, f'A{rt}', '合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'F{rt}', f'=SUM(F{R0}:F{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    for C in (PL_C1, PL_C2):
        for k in PL_BLK:
            if k == '净利率':
                put(ws, f'{C[k]}{rt}', f'=IF(N({C["收入"]}{rt})=0,"",{C["净利"]}{rt}/{C["收入"]}{rt})', F_AUTOB, FILL_TOT, PCT, AC)
            else:
                put(ws, f'{C[k]}{rt}', f'=SUM({C[k]}{R0}:{C[k]}{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'{PL_TIP}{rt}', None, fill_=FILL_TOT)

    # ── 跟公司利润表对账（本期）
    RB = PL_RB
    section(ws, RB, 'A', 'Q', '跟公司利润表对账（本期：起～止）：项目净利加减下面几项＝公司净利润', C_PROJ)
    dr_s, dr_f, dr_fk, dr_y, dr_k = _dr('收', X, Y), _dr('付', X, Y), _dr('付', X, Y, '开票日期'), _dr('应', X, Y), _drk(X, Y)
    tans = '+'.join(f'SUMIFS(考_摊{j + 1},{dr_k})' for j in range(NPAIR))
    xa = f'({X}-1)'
    hidc = [  # 隐藏：AV 名称、AW 值（第 RB+1 行起）
        ('没项目开票止', f'=SUMIFS(应_开票额,应_日期,"<="&{Y})-SUM($AI${R0}:$AI${R1})'),
        ('没项目收款止', f'=SUMIFS(收_净额,收_归类,"工程款收款",收_日期,"<="&{Y})-SUM($AJ${R0}:$AJ${R1})'),
        ('没项目产值止', f'=SUMIFS(应_产值额,应_日期,"<="&{Y})-SUM($AK${R0}:$AK${R1})'),
        ('没项目开票起前', f'=SUMIFS(应_开票额,应_日期,"<="&{xa})-SUM($AL${R0}:$AL${R1})'),
        ('没项目收款起前', f'=SUMIFS(收_净额,收_归类,"工程款收款",收_日期,"<="&{xa})-SUM($AM${R0}:$AM${R1})'),
        ('没项目产值起前', f'=SUMIFS(应_产值额,应_日期,"<="&{xa})-SUM($AN${R0}:$AN${R1})'),
        ('没项目收入', None),
        ('待摊池', f'=-SUMIFS(收_净额,收_归属,"待摊费用",{dr_s})+SUMIFS(考_未分摊,考_计提,1,{dr_k})'),
        ('已摊', f'=SUMIFS(收_摊费,{dr_s})+SUMIFS(付_摊费,{dr_f})+{tans}'),
        # 公司利润表（§9）另算一遍
        ('工程收入', None),
        ('其他收入', f'=SUMIFS(收_净额,收_归类,"其他收入",{dr_s})'),
        ('材料', f'=SUMIFS(付_应付额,付_类型,"材料",{dr_f})-SUMIFS(收_净额,收_成本类,"材料",收_归属,"项目成本",{dr_s})'),
        ('分包', f'=SUMIFS(付_应付额,付_类型,"分包",{dr_f})-SUMIFS(收_净额,收_成本类,"分包",收_归属,"项目成本",{dr_s})'),
        ('机械', f'=SUMIFS(付_应付额,付_类型,"机械",{dr_f})-SUMIFS(收_净额,收_成本类,"机械",收_归属,"项目成本",{dr_s})'),
        ('其他直接费', f'=SUMIFS(付_应付额,{dr_f})-SUMIFS(付_应付额,付_类型,"材料",{dr_f})-SUMIFS(付_应付额,付_类型,"分包",{dr_f})'
                    f'-SUMIFS(付_应付额,付_类型,"机械",{dr_f})-SUMIFS(收_净额,收_成本类,"其他直接费",收_归属,"项目成本",{dr_s})'
                    f'-SUMIFS(收_净额,收_成本类,"管理费用",收_归属,"项目成本",{dr_s})'),
        ('人工', f'=SUMIFS(考_应发,考_计提,1,{dr_k})-SUMIFS(考_未分摊,考_计提,1,{dr_k})'
               f'-SUMIFS(收_净额,收_成本类,"人工",收_归属,"项目成本",{dr_s})'),
        ('管理费用', f'=-SUMIFS(收_净额,收_归属,"待摊费用",{dr_s})-(SUMIFS(收_净额,收_归属,"公司费用",{dr_s})'
                 f'-SUMIFS(收_净额,收_成本类,"财务费用",{dr_s})-SUMIFS(收_净额,收_成本类,"营业外支出",{dr_s})'
                 f'-SUMIFS(收_净额,收_成本类,"未分类",{dr_s}))+SUMIFS(考_未分摊,考_计提,1,{dr_k})'),
        ('财务费用', f'=-SUMIFS(收_净额,收_成本类,"财务费用",{dr_s})'),
        ('营业外支出', f'=-SUMIFS(收_净额,收_成本类,"营业外支出",{dr_s})'),
        ('未分类', f'=SUMIFS(收_净额,收_成本类,"未分类",{dr_s})'),
        ('税金', f'=(SUMIFS(应_销项税,{dr_y})-SUMIFS(付_进项税,{dr_fk}))*(1+P_附加税率)'),
        ('净利润', None),
    ]
    HC = {k: f'$AW${RB + 1 + i}' for i, (k, _f) in enumerate(hidc)}
    for i, (k, f) in enumerate(hidc):
        r = RB + 1 + i
        _hid(ws, f'AV{r}', k)
        if k == '没项目收入':
            f = (f'=({rev_formula(HC["没项目开票止"], HC["没项目收款止"], HC["没项目产值止"])})'
                 f'-({rev_formula(HC["没项目开票起前"], HC["没项目收款起前"], HC["没项目产值起前"])})')
        elif k == '工程收入':
            f = f'=${PL_C1["收入"]}${rt}+{HC["没项目收入"]}'
        elif k == '净利润':
            f = (f'={HC["工程收入"]}+{HC["其他收入"]}+{HC["未分类"]}-{HC["材料"]}-{HC["分包"]}-{HC["机械"]}-{HC["人工"]}'
                 f'-{HC["其他直接费"]}-{HC["管理费用"]}-{HC["财务费用"]}-{HC["营业外支出"]}-{HC["税金"]}')
        _hid(ws, f'AW{r}', f)
    header(ws, RB + 1, [('A', ''), ('B', '项目'), ('F', '金额'), ('G', '说明')], C_PROJ, height=24)
    ws.merge_cells(f'B{RB + 1}:E{RB + 1}')
    ws.merge_cells(f'G{RB + 1}:Q{RB + 1}')
    lines = [
        ('各项目净利合计（上面合计行「本期 · 净利」）', f'=${PL_C1["净利"]}${rt}', '每个工程项目：确认收入−直接成本−税金−分摊管理费'),
        ('减：没选项目的应付登记（算公司成本，没进任何项目）', f'=-(SUMIFS(付_应付额,{dr_f})-SUM($AR${R0}:$AR${R1}))',
         '在【应付登记】给这些行选上项目，就会进那个项目的成本'),
        ('加：上面那些应付里专票的进项税（公司少交的税）', f'=(SUMIFS(付_进项税,{dr_fk})-SUM($AS${R0}:$AS${R1}))*(1+P_附加税率)', ''),
        ('减：待摊费用里没摊出去的', f'=-({HC["待摊池"]}-{HC["已摊"]})',
         f'="这段时间待摊 "&TEXT({HC["待摊池"]},"#,##0.00")&"，摊到项目 "&TEXT({HC["已摊"]},"#,##0.00")&"（某一年没有施工费时摊不出去）"'),
        ('减：公司费用（贷款利息、手续费、罚款、没认出来的收支、归属选了「公司费用」的）', f'=SUMIFS(收_净额,收_归属,"公司费用",{dr_s})',
         '这些不摊给项目，直接算公司的'),
        ('加：其他收入（利息收入等）', f'=SUMIFS(收_净额,收_归类,"其他收入",{dr_s})', ''),
        ('加：没选项目的工程款（按收入口径算的收入）', f'={HC["没项目收入"]}', '【收支登记】工程款没选工程项目的'),
    ]
    r = RB + 2
    first = r
    for lab, f, note in lines:
        ws.merge_cells(f'B{r}:E{r}')
        ws.merge_cells(f'G{r}:Q{r}')
        put(ws, f'B{r}', lab, F_TXT, align=AL)
        put(ws, f'F{r}', f, F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'G{r}', note, F_NOTE, align=AL)
        r += 1
    tot = r
    for lab, f, note, fl in (
            ('＝ 公司净利润（照上面加减）', f'=SUM(F{first}:F{tot - 1})', '', FILL_SUB),
            ('公司净利润（按【利润表】的算法另算一遍）', f'={HC["净利润"]}',
             (f'="【利润表】现在的期间："&TEXT(汇_利润起,"yyyy/mm/dd")&"～"&TEXT(汇_利润止,"yyyy/mm/dd")'
              f'&IF(AND(汇_利润起={X},汇_利润止={Y}),"，跟这里一样，利润表的净利润应该等于这个数","，跟这里不一样（把两边起止填成一样再比）")'), FILL_SUB),
            ('差额（应为 0）', f'=ROUND(F{tot}-F{tot + 1},2)',
             f'=IF(F{tot + 2}=0,"✓ 对上了","✗ 对不上：工程项目超过 {NP} 个没列全，或者有数据问题（看【数据校验】）")', FILL_TOT)):
        ws.merge_cells(f'B{r}:E{r}')
        ws.merge_cells(f'G{r}:Q{r}')
        put(ws, f'B{r}', lab, F_TXTB, fl, align=AL)
        put(ws, f'F{r}', f, F_AUTOB, fl, MONEY, AR)
        put(ws, f'G{r}', note, F_RED if lab.startswith('差额') else F_NOTE, align=AL)
        r += 1
    hide(ws, *[CL(i) for i in range(CI('AD'), CI('AW') + 1)])
    ws.freeze_panes = f'C{R0}'
    print_setup(ws, '5:6', landscape=True)
    return ws


# ═══════════════════════════ 费用分摊 ═══════════════════════════
AL_Y0 = 17                        # ① 年度表：第 17～24 行，25 合计
AL_P0, AL_NP = 30, 30             # ② 所选年度各项目：第 30 行起 30 行，第 29 行合计
AL_D0 = AL_P0 + AL_NP + 4         # ③ 待摊明细：数据第一行
AL_ND, AL_NK = 500, 200           # ③ 流水最多 500 条、考勤最多 200 条
ALC = dict(Y='$P$3', D0='$P$4', D1='$P$5', YI='$P$6', RATE='$P$7')


def build_alloc(wb, ctx):
    ws = wb[SH_ALLOC]
    W = {'A': 11, 'B': 16, 'C': 14, 'D': 14, 'E': 14, 'F': 12, 'G': 14, 'H': 14, 'I': 30, 'J': 2, 'K': 11, 'L': 10, 'M': 14}
    widths(ws, W)
    tip = ('💡 黄格选年度（空着＝首页报表年度）。全自动：哪些开支要摊，在【收支登记】O 列「费用归属」选；'
           '这张表算每年的费率、每个项目摊多少，下面列出所选年度的每一笔待摊。')
    title(ws, '费用分摊（办公室开支、管理人员工资怎么摊到各项目）', 'M', C_PROJ, tip=tip)
    Y, D0, D1, YI, RATE = (ALC[k] for k in ('Y', 'D0', 'D1', 'YI', 'RATE'))
    for i, (k, f) in enumerate([('年度', '=IF(ISNUMBER($C$3),INT($C$3),P_年度)'), ('年初', f'=DATE({Y},1,1)'), ('年末', f'=DATE({Y},12,31)'),
                                ('第几年', f'={Y}-{YEAR0}+1'), ('费率', f'=IFERROR(INDEX(汇_费率,{YI}),0)')]):
        _hid(ws, f'O{3 + i}', k)
        _hid(ws, f'P{3 + i}', f)
    selector(ws, 'B3', '年度', 'C3', None, '"' + ','.join(str(YEAR0 + i) for i in range(NYEARS)) + '"', fmt='0',
             prompt='选年度；空着＝首页报表年度')
    ws.merge_cells('D3:I3')
    put(ws, 'D3', (f'="实际看："&{Y}&" 年"&IF(ISNUMBER($C$3),"","（空着＝首页报表年度）")&"；分摊依据「"&P_分摊依据&"」，这一年费率 "'
                   f'&TEXT({RATE},"0.00%")&IF(OR({Y}<{YEAR0},{Y}>{YEAR0 + NYEARS - 1}),"　⚠ 只有 {YEAR0}～{YEAR0 + NYEARS - 1} 年","")'),
        F_NOTE, align=AL, border=False)
    home_link(ws, 'M3')

    # ── 说明（大白话）
    section(ws, 5, 'A', 'M', '说明：什么是待摊、怎么选、按什么摊、摊到什么时候', C_PROJ)
    notes = [
        '① 什么是待摊：公司每天都有不属于哪一个工地的开支——办公室房租水电、办公用品、招待、车辆油费、管理人员在办公室那几天的工资等。'
        '这些先放进「待摊池」，按各项目干了多少活分下去，项目利润才算得全。',
        '② 怎么选：在【收支登记】O 列「费用归属」选——项目成本（只算这一个项目，要选项目名称）、待摊费用（进待摊池，分给所有项目）、公司费用（留在公司，不摊）。'
        '空着＝自动：选了工程项目的材料、人工、分包、机械、其他→项目成本；办公室、管理费用类→待摊；贷款利息、手续费、罚款→公司费用。',
        '="③ 按什么摊：【基本信息】①「管理费分摊依据」现在是「"&P_分摊依据&"」"&IF(P_分摊依据="直接成本","（材料＋人工＋分包＋机械＋其他直接费）","（人工＋分包＋机械，一般建筑安装用这个）")'
        '&"。工人、分包、机械用得越多，越费管理，就多摊；材料多少跟管理关系不大。"',
        '④ 一年一个费率，不用填：费率＝这一年的待摊合计 ÷ 这一年所有项目的分摊基数；项目这一年摊到＝它这一年的基数 × 费率。'
        '同一年的待摊正好摊完，一分不多一分不少；年中看到的是「到目前为止」的费率，到年底自然定下来。',
        '⑤ 摊到什么时候：不用填截止日期。项目完工后没有人工、分包、机械（基数＝0），自然就不再摊；结算、收尾款、质保金都不影响。'
        '哪个月又有人去干活（比如保修），那个月就照样摊一点。哪一年一个项目都没有施工费，那年的待摊摊不出去，留在公司（下面「没摊出去」）。',
        '⑥ 公司费用不摊：贷款利息、银行手续费、罚款滞纳金这些跟干多少活无关，直接进公司利润表，不分给项目。',
        '⑦ 明面上是货款、其实是红包/好处费的：不要记成材料款（会算进材料成本，填了供应商还会冲掉他的应付）。收支项目选「业务招待费」这类管理费用，摘要写清楚，'
        '费用归属选「待摊费用」，就进待摊池按年分摊；如果确定是为某一个项目花的，就选那个项目、归属选「项目成本」，算它的其他直接费。',
        '⑧ 管理人员的工资：【考勤工资】里填了某个工地几天的，那几天的工资直接算那个项目的人工；在办公室（公司管理）的天数、没对上项目的部分进待摊。',
    ]
    _note_rows(ws, 6, 'A', 'M', notes, height=32)

    # ── ① 年度表
    section(ws, AL_Y0 - 2, 'A', 'I', '① 每年的待摊和费率', C_PROJ)
    header(ws, AL_Y0 - 1, [('A', '年'), ('B', '待摊（收支登记）'), ('C', '管理人员工资\n（考勤没分到项目的）'), ('D', '待摊合计'),
                           ('E', '分摊基数\n（各项目的施工费）'), ('F', '费率'), ('G', '已摊到项目'), ('H', '没摊出去'), ('I', '说明')],
           C_PROJ, height=44)
    ycols = [('B', '汇_待摊池收'), ('C', '汇_待摊池考'), ('D', '汇_待摊池'), ('E', '汇_基数'), ('F', '汇_费率'), ('G', '汇_已摊'), ('H', '汇_未摊')]
    for i in range(NYEARS):
        r = AL_Y0 + i
        put(ws, f'A{r}', f'=INDEX(汇_年,{i + 1})', F_TXTB, align=AC)
        for c, nm in ycols:
            put(ws, f'{c}{r}', f'=INDEX({nm},{i + 1})', F_AUTOB if c in ('D', 'F') else F_AUTO, fmt=PCT if c == 'F' else MONEY,
                align=AC if c == 'F' else AR)
        put(ws, f'I{r}', f'=IF(AND(ROUND(D{r},2)<>0,ROUND(E{r},2)=0),"这一年没有施工费，待摊摊不出去，留在公司",'
                         f'IF(ROUND(H{r},2)<>0,"还有没摊出去的",""))', F_NOTE, align=AL)
    rt = AL_Y0 + NYEARS
    put(ws, f'A{rt}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c, _nm in ycols:
        if c == 'F':
            put(ws, f'F{rt}', None, fill_=FILL_TOT)
        else:
            put(ws, f'{c}{rt}', f'=SUM({c}{AL_Y0}:{c}{rt - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'I{rt}', None, fill_=FILL_TOT)
    ws.conditional_formatting.add(f'A{AL_Y0}:I{rt - 1}', FormulaRule(formula=[f'$A{AL_Y0}={Y}'], fill=FILL_HI, font=F_TXTB))

    # ── ② 所选年度各项目（只列这一年有基数或摊费的工程项目）
    # 隐藏：Q 第几条（常数，第 2 行起）；R 项目名、S 转义、T 基数、U 摊费（第 2～201 行）；V 计数（counter）
    for i in range(N_CASH):
        _hid(ws, f'Q{2 + i}', i + 1)
    for i in range(N_PJ):
        r = 2 + i
        n, nm, pe = f'$Q{r}', f'$R{r}', f'$S{r}'
        _hid(ws, f'R{r}', f'=IF(INDEX(项目_类型,{n})="工程",INDEX(项目_名称,{n})&"","")')
        _hid(ws, f'S{r}', f'=IF({nm}="","",{esc(nm)})')
        _hid(ws, f'T{r}', (f'=IF({nm}="",0,SUMIFS(收_基数,收_项目,{pe},{_dr("收", D0, D1)})+SUMIFS(付_基数,付_项目,{pe},{_dr("付", D0, D1)})+'
                           + _kq(pe, D0, D1) + ')'))
        _hid(ws, f'U{r}', (f'=IF({nm}="",0,SUMIFS(收_摊费,收_项目,{pe},{_dr("收", D0, D1)})+SUMIFS(付_摊费,付_项目,{pe},{_dr("付", D0, D1)})+'
                           + _kq(pe, D0, D1, '摊', jt=False) + ')'))
    counter(ws, 'V', 2, N_PJ, lambda i: f'AND($R{2 + i}<>"",OR(ROUND($T{2 + i},2)<>0,ROUND($U{2 + i},2)<>0))')
    NPC = cnt('V', 2, N_PJ)
    section(ws, AL_P0 - 3, 'A', 'I', f'="② "&{Y}&" 年：各项目摊到多少（只列这一年有施工费的工程项目）"', C_PROJ)
    header(ws, AL_P0 - 2, [('A', '序号'), ('B', '项目'), ('C', '状态'), ('D', '完工'), ('E', '分摊基数\n（这一年）'), ('F', '费率'),
                           ('G', '摊到的管理费'), ('H', '占比'), ('I', '说明')], C_PROJ, height=36)
    c = put(ws, f'K{AL_P0 - 3}', f'↓ 待摊明细（第 {AL_D0 - 3} 行）', Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'),
            align=AL, border=False)
    link(c, SH_ALLOC, f'A{AL_D0 - 3}')
    rt2 = AL_P0 - 1
    P1 = AL_P0 + AL_NP - 1
    put(ws, f'A{rt2}', '合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'B{rt2}', f'="共 "&{NPC}&" 个项目"&IF({NPC}>{AL_NP},"（只列前 {AL_NP} 个）","")', F_NOTE, FILL_TOT, align=AL)
    put(ws, f'C{rt2}', None, fill_=FILL_TOT)
    put(ws, f'D{rt2}', None, fill_=FILL_TOT)
    put(ws, f'E{rt2}', f'=SUM(E{AL_P0}:E{P1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'F{rt2}', f'={RATE}', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, f'G{rt2}', f'=SUM(G{AL_P0}:G{P1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'H{rt2}', f'=IF(G{rt2}=0,"",1)', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, f'I{rt2}', (f'=IF(ROUND(G{rt2}-IFERROR(INDEX(汇_已摊,{YI}),0),2)=0,"✓ 等于年度表「已摊到项目」",'
                        f'"✗ 跟年度表「已摊到项目」差 "&TEXT(G{rt2}-IFERROR(INDEX(汇_已摊,{YI}),0),"#,##0.00"))'), F_RED, FILL_TOT, align=AL)
    for i in range(AL_NP):
        r = AL_P0 + i
        ix = f'$X{r}'
        _hid(ws, f'W{r}', i + 1)
        _hid(ws, f'X{r}', '=' + kth(f'$W{r}', 'V', 2, N_PJ))
        ws[f'A{r}'] = f'=IF({ix}=0,"",$W{r})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX($R$2:$R${N_PJ + 1},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX(项目_状态,{ix})&"")'
        ws[f'D{r}'] = f'=IF({ix}=0,"",IF(INDEX(项目_完工,{ix})>0,INDEX(项目_完工,{ix}),""))'
        ws[f'E{r}'] = f'=IF({ix}=0,"",INDEX($T$2:$T${N_PJ + 1},{ix}))'
        ws[f'F{r}'] = f'=IF({ix}=0,"",{RATE})'
        ws[f'G{r}'] = f'=IF({ix}=0,"",INDEX($U$2:$U${N_PJ + 1},{ix}))'
        ws[f'H{r}'] = f'=IF(OR({ix}=0,N($G${rt2})=0),"",G{r}/$G${rt2})'
        ws[f'I{r}'] = (f'=IF({ix}=0,"",IF(AND(INDEX(项目_完工,{ix})>0,INDEX(项目_完工,{ix})<={D1}),'
                       f'"完工日以后没有施工费就不再摊",""))')
    _list_area(ws, AL_P0, AL_NP, list('ABCDEFGHI'), {'D': DATE, 'E': MONEY, 'F': PCT, 'G': MONEY, 'H': PCT},
               {'B': AL, 'E': AR, 'G': AR, 'I': AL})
    for r in range(AL_P0, AL_P0 + AL_NP):
        ws[f'B{r}'].font = F_TXTB
        ws[f'G{r}'].font = F_AUTOB

    # ── ③ 待摊明细
    D0r = AL_D0
    section(ws, D0r - 3, 'A', 'I', f'="③ "&{Y}&" 年待摊明细：收支登记里归属＝待摊费用的每一笔"', C_CASH)
    section(ws, D0r - 3, 'K', 'M', '管理人员工资（考勤没分到项目的）', C_HOME)
    # 隐藏：Y 收支排序键（第 2～5001 行）、Z 考勤排序键（第 2～2001 行），第 1 行放个数；AA/AC 第 k 条是第几条，AB k
    for i in range(N_CASH):
        r = 2 + i
        n = f'$Q{r}'
        _hid(ws, f'Y{r}', (f'=IF(OR(INDEX(收_归属,{n})<>"待摊费用",INDEX(收_日期,{n})<{D0},INDEX(收_日期,{n})>{D1}),"",'
                           f'INDEX(收_日期,{n})*10000+{n})'))
    for i in range(N_ATT):
        r = 2 + i
        n = f'$Q{r}'
        _hid(ws, f'Z{r}', (f'=IF(OR(INDEX(考_计提,{n})<>1,ROUND(INDEX(考_未分摊,{n}),2)=0,INDEX(考_年月,{n})<{Y}*100+1,'
                           f'INDEX(考_年月,{n})>{Y}*100+12),"",INDEX(考_月,{n})*10000+{n})'))
    _hid(ws, 'Y1', f'=COUNT($Y$2:$Y${N_CASH + 1})')
    _hid(ws, 'Z1', f'=COUNT($Z$2:$Z${N_ATT + 1})')
    KY, KZ = f'$Y$2:$Y${N_CASH + 1}', f'$Z$2:$Z${N_ATT + 1}'
    header(ws, D0r - 1, [('A', '日期'), ('B', '账户'), ('C', '收支项目'), ('D', '项目名称\n（录的）'), ('E', '金额'),
                         ('F', '归属\n手选/自动'), ('G', '摘要')], C_CASH, height=36)
    ws.merge_cells(f'G{D0r - 1}:I{D0r - 1}')
    header(ws, D0r - 1, [('K', '月份'), ('L', '姓名'), ('M', '金额')], C_HOME, height=36)
    ws.merge_cells(f'A{D0r - 2}:D{D0r - 2}')
    put(ws, f'A{D0r - 2}', f'="共 "&$Y$1&" 笔"&IF($Y$1>{AL_ND},"，只列前 {AL_ND} 笔","")&"　合计 →"', F_NOTE, FILL_TOT, align=AR)
    D1r = D0r + AL_ND - 1
    put(ws, f'E{D0r - 2}', f'=SUM(E{D0r}:E{D1r})', F_AUTOB, FILL_TOT, MONEY, AR)
    ws.merge_cells(f'F{D0r - 2}:I{D0r - 2}')
    put(ws, f'F{D0r - 2}', (f'=IF(ROUND(E{D0r - 2}-IFERROR(INDEX(汇_待摊池收,{YI}),0),2)=0,"✓ 等于年度表「待摊（收支登记）」",'
                            f'"✗ 跟年度表差 "&TEXT(E{D0r - 2}-IFERROR(INDEX(汇_待摊池收,{YI}),0),"#,##0.00"))'), F_RED, FILL_TOT, align=AL)
    put(ws, f'K{D0r - 2}', f'="共 "&$Z$1&" 条"', F_NOTE, FILL_TOT, align=AC)
    put(ws, f'L{D0r - 2}', (f'=IF(ROUND(M{D0r - 2}-IFERROR(INDEX(汇_待摊池考,{YI}),0),2)=0,"✓","✗")'), F_RED, FILL_TOT, align=AC)
    K1r = D0r + AL_NK - 1
    put(ws, f'M{D0r - 2}', f'=SUM(M{D0r}:M{K1r})', F_AUTOB, FILL_TOT, MONEY, AR)
    for i in range(AL_ND):
        r = D0r + i
        k, ix = f'$AB{r}', f'$AA{r}'
        _hid(ws, f'AB{r}', i + 1)
        _hid(ws, f'AA{r}', f'=IF({k}>$Y$1,0,MOD(SMALL({KY},{k}),10000))')
        ws[f'A{r}'] = f'=IF({ix}=0,"",INDEX(收_日期,{ix}))'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX(收_账户,{ix})&"")'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX(收_收支项目,{ix})&"")'
        ws[f'D{r}'] = f'=IF({ix}=0,"",INDEX(收_项目原,{ix})&"")'
        ws[f'E{r}'] = f'=IF({ix}=0,"",-INDEX(收_净额,{ix}))'
        ws[f'F{r}'] = (f'=IF({ix}=0,"",IF(INDEX(收_归属原,{ix})="待摊费用","手选",IF(AND(INDEX(收_成本类,{ix})<>"管理费用",'
                       f'INDEX(收_项目,{ix})=""),"自动（没选工程项目）","自动")))')
        ws[f'G{r}'] = f'=IF({ix}=0,"",INDEX(收_摘要,{ix})&"")'
        if i < AL_NK:
            kk, jx = f'$AB{r}', f'$AC{r}'
            _hid(ws, f'AC{r}', f'=IF({kk}>$Z$1,0,MOD(SMALL({KZ},{kk}),10000))')
            ws[f'K{r}'] = f'=IF({jx}=0,"",INDEX(考_月,{jx}))'
            ws[f'L{r}'] = f'=IF({jx}=0,"",INDEX(考_姓名,{jx})&"")'
            ws[f'M{r}'] = f'=IF({jx}=0,"",INDEX(考_未分摊,{jx}))'
    _list_area(ws, D0r, AL_ND, list('ABCDEFG'), {'A': DATE, 'E': MONEY}, {'B': AL, 'C': AL, 'D': AL, 'E': AR, 'G': AL})
    _list_area(ws, D0r, AL_NK, ['K', 'L', 'M'], {'K': MONTHF, 'M': MONEY}, {'M': AR})
    hide(ws, *[CL(i) for i in range(CI('N'), CI('AC') + 1)])
    ws.freeze_panes = 'A4'
    print_setup(ws, '1:3', landscape=True)
    return ws


def build(wb, ctx=None):
    build_alloc(wb, ctx)
    build_pl(wb, ctx)
    build_book(wb, ctx)
