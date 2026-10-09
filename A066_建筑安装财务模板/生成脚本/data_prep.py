# -*- coding: utf-8 -*-
"""A066 第二版 · 数据导入。

build_ctx() 读 ../参考/1008上传/ 里用户的原始文件（openpyxl data_only=True，只读数据，不执行任何内容），
整理成生成账套用的 ctx（键见 设计.md / layout.py）：params accounts inc_items exp_items projects customers
suppliers persons rates opening cash ar ap att inv arrears log questions checks。

直接运行：python3 -I data_prep.py → 打印摘要、核对结果，并在本目录写 _ctx.json（含身份证/银行卡，别外传）。
说明文档：导入说明.md。
"""
import os
import re
import sys
import json
import zipfile
import warnings
import datetime as dt
import itertools
import collections
import xml.etree.ElementTree as ET

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)
import layout as L  # noqa: E402

warnings.filterwarnings('ignore', module='openpyxl')
import openpyxl  # noqa: E402

SRC = os.path.normpath(os.path.join(HERE, '..', '参考', '1008上传'))
F_REG, F_AR, F_MAT, F_FB, F_JX = '收支登记表.xlsx', '应收账款.xlsx', '应付账款_材料.xlsx', '应付账款_分包.xlsx', '机械费.xlsx'
F_ARREAR, F_U1007, F_NH, F_PERS = '2025未发清工资.xlsx', 'A066_用户1007版.xlsx', '聂辉报销明细.xlsx', '个人借款及个人收支明细.xlsx'
CUTOFF = dt.date(2026, 9, 30)
BANK_EXPECT = {'九江银行': 2729.48, '中国农业银行': 92.04, '中国工商银行': 147.99, '现金': 3060.0, '农民工专户': 0.0, '九安账户': 0.0}


# ═════════════════════════════ 小工具 ═════════════════════════════
def P(fn):
    return os.path.join(SRC, fn)


def todate(v):
    """日期、Excel 序列号、'2025-01-02' 文本 → date；认不出 → None"""
    if v is None or v == '':
        return None
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v)) if 1 <= v < 80000 else None
    if isinstance(v, str):
        m = re.match(r'^\s*(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})', v)
        if m:
            try:
                return dt.date(int(m[1]), int(m[2]), int(m[3]))
            except ValueError:
                return None
    return None


def okdate(d):
    return d is not None and d.year >= 2000


def num(v):
    if v is None or v == '' or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(',', '')
        try:
            return float(s)
        except ValueError:
            return None
    return None


def r2(x):
    v = round((x or 0) + 0.0, 2)
    return v if v != 0 else 0.0


def s(v):
    return '' if v is None else str(v).strip()


def money(x):
    return f'{r2(x):,.2f}'


def hnorm(v):
    return re.sub(r'\s+', '', s(v))


def joinn(*parts):
    return '；'.join(p for p in parts if p)


def cjk(t):
    return bool(re.fullmatch(r'[一-龥]{2,3}', t or ''))


# ═════════════════════════════ 名称对照（固定表） ═════════════════════════════
# 项目：各表写法 → 项目简称（简称取应收/应付分表、1007 项目档案的叫法）
PJ_ALIAS = {
    '德安县养老中心': '德安养老中心', '德安县科创中心': '科创中心', '九安架空层': '九安架空层',
    '德安县数字产业园': '数字产业园', '数字经济产业园': '数字产业园', '德安县邹桥派出所': '邹桥派出所',
    '抚州资溪养老基地1号楼': '抚州1号楼', '抚州资溪养老基地8号楼': '抚州8号楼', '抚州资溪养老基地2号楼': '抚州2号楼',
    '德安总部共青农商行': '德安农商行总部', '江西赣粤高速公路工程有限责任公司': '赣粤高速',
    '温州盛为装饰工程有限公司': '温州盛为', '九江欣埠置业有限公司': '九江欣埠置业',
    '隆世光伏/共青办公楼/德祥工地': '隆世光伏', '公司管理': '办公室',
}
PSEUDO_PJ = {'个人借款': '借款类收支项目，不是工程', '贷款': '银行贷款，不是工程', '姚俊强': '过账人员，改记在「人员」'}
COMPANY_PJ = '办公室'

# 甲方：写法 → 规范名
CUS_LIST = [  # (规范名, 简称/别名)
    ('德安县博河产业控股集团有限公司', '博河'),
    ('九江浔辉建筑工程有限公司', '浔辉'),
    ('江西揽之建设工程有限公司', '揽之'),
    ('德安县市政工程公司', '市政'),
    ('德安县水利水电建筑工程公司', '水利水电、德安县水利水电建筑工程有限公司'),
    ('江西中寰建筑装饰有限公司', '中寰'),
    ('共青农村商业银行股份有限公司', '共青农商行、共青农商银行、德安农商银行、农村商业银行股份有限公司德安县总部'),
    ('江西赣粤高速公路工程有限责任公司', '赣粤高速、江西赣粤高速、江西赣粤高速公路'),
    ('联通（江西）产业互联网有限公司', '联通'),
    ('温州盛为装饰工程有限公司', '温州盛为'),
    ('九江欣埠置业有限公司', '欣埠'),
]
CUS_ALIAS = {'德安县水利水电建筑工程有限公司': '德安县水利水电建筑工程公司', '共青农商行': '共青农村商业银行股份有限公司',
             '共青农商银行': '共青农村商业银行股份有限公司', '德安农商银行': '共青农村商业银行股份有限公司',
             '江西赣粤高速': '江西赣粤高速公路工程有限责任公司', '江西赣粤高速公路': '江西赣粤高速公路工程有限责任公司'}

# 应付分表供应商：摘要/收款方关键词（付款配对用）
MAT_KW = {
    '筑强': ['筑强', '绿谷嘉园', '福晟官邸'], '九江安洋钢材熊总': ['安洋'], '鑫城五金庄总': ['鑫城', '安迪', '李美坚'],
    '浔阳楼电缆': ['浔阳楼', '日昌', '鄂缆', '坚宝', '双帝'], '启航五金': ['启航'], '杭州行消物资贸易': ['行消'],
    '南昌盈厚建材': ['盈厚'], '浙缆元通电缆': ['浙缆', '元通'], '江西允诺消防': ['允诺'], '江西沐丰水箱': ['沐丰'],
    '上海诚果泵阀': ['诚果'], '佩力通风': ['佩力'], '浙江寅盾智能': ['寅盾'], '九江中和安': ['中和安'], '江西弘文消防': ['弘文'],
    '华翔电缆桥架': ['华翔'], '温州冠林物资': ['冠林'], '温州创泰电气': ['创泰'], '浙江杭风通风': ['杭风'], '浙江尚杰建材': ['尚杰'],
    '江西跃池建材': ['跃池'], '天闽消防': ['天闽'], '瑞安中拓拱水': ['中拓'], '辰南贸易': ['辰南'], '南昌航卓建材': ['航卓'],
    '德安远达建材': ['远达'], '九江恒讯消防': ['恒讯'], '四川久远': ['久远'], '温州爱思强贸易': ['爱思强贸易'],
    '零星供应商': ['众贝源', '共安机电', '立邦', '春虹', '安南五金', '敏芬', '伦川', '孙法炳'],
}
SUB_KW = {
    '恩旗老叶': ['恩旗', '老叶'], '吴群': ['吴群'], '吴聪': ['吴聪'], '九江朱总': ['朱总'], '共安门业': ['共安门业'],
    '图揽成善': ['图揽', '成善'], '祝冬冬': ['祝冬冬'], '张小朋': ['张小朋'], '徐怀生': ['徐怀生'], '龚文坚': ['龚文坚'],
    '程章赐': ['程章赐'],
}
SUB_EXTRA = {'严程': ['严程装'], '刘剑': ['刘剑分包']}          # 登记表里有付款、分包表里没有的两家
JX_KW = {'杨工轮挖': ['杨工', '翌年行'], '徐工轮挖': ['徐工'], '德安县恒建工程部': ['恒建', '舒文龙'], '航帆增辉': ['航帆'],
         '章德吉货运': ['章德吉']}
JX_NAME = {'杨工': '杨工轮挖', '徐工': '徐工轮挖'}

# 同一家公司的其他写法（登记表「供应商名称」→ 供应商规范名）。没配上付款的行也按这个归。
SAME = {
    '江西筑强建材有限公司': '筑强', '九江安洋钢材销售有限公司': '九江安洋钢材熊总', '德安县鑫城五金店': '鑫城五金庄总',
    '德安县启航五金机电超市': '启航五金', '杭州行消物资贸易有限公司': '杭州行消物资贸易', '南昌盈厚建材有限公司': '南昌盈厚建材',
    '浙缆元通电缆控股有限公司': '浙缆元通电缆', '江西允诺消防设备有限公司': '江西允诺消防', '江西沐丰供水设备有限公司': '江西沐丰水箱',
    '上海诚果泵阀制造有限公司永嘉分公司': '上海诚果泵阀', '江西佩力通风设备有限公司': '佩力通风',
    '浙江寅盾智能消防有限公司': '浙江寅盾智能', '九江市中和安消防科技有限公司': '九江中和安', '江西弘文消防科技有限公司': '江西弘文消防',
    '江西华翔节能科技有限公司': '华翔电缆桥架', '温州冠林物资有限公司': '温州冠林物资', '温州创泰电气设备有限公司': '温州创泰电气',
    '浙江杭风通风设备有限公司': '浙江杭风通风', '浙江尚杰建材有限公司': '浙江尚杰建材', '江西跃池建材有限公司': '江西跃池建材',
    '江西天闽消防科技有限公司': '天闽消防', '瑞安市中拓供水设备有限公司': '瑞安中拓拱水', '江西省辰南贸易有限公司': '辰南贸易',
    '南昌航卓建材装鉓有限公司': '南昌航卓建材', '德安县远达建材店': '德安远达建材', '九江恒讯消防安装工程有限公司': '九江恒讯消防',
    '四川久远智能消防设备有限公司': '四川久远', '四川久远智能消防设备有限责任公司': '四川久远',
    '九江恩旗暖通设备有限公司': '恩旗老叶', '江西共安门业有限公司': '共安门业', '江西图揽金属科技有限公司': '图揽成善',
    '江西成善科技有限公司': '图揽成善', '朱总': '九江朱总', '吴群': '吴群', '吴聪': '吴聪', '龚文坚': '龚文坚', '徐怀生': '徐怀生',
    '九江翌年行建筑工程劳务有限公司': '杨工轮挖', '舒文龙': '德安县恒建工程部',
    '九江财小保企业服务服限公司': '九江财小保企业服务有限公司', '中国平安保险股份有限公司': '中国平安财产保险股份有限公司广东分公司',
}
MEMO_SUP = {'迈腾建筑工程有限公司': '九江迈腾建筑工程有限公司', '柯源建材': '江西柯源建材有限公司'}
NON_PERSON_MEMO = ['迈腾建筑工程有限公司', '柯源建材', '瑞金项目劳务费']
NAME_STOP = {'开孔', '油漆工', '工资', '奖金', '退回', '补贴', '代缴', '员工', '个税'}

# 人名错别字：(错, 对) —— 对的那个以 1007 考勤/工资标准/往来单位里出现的为准（运行时核对次数）
NAME_TYPO = [('何正佣', '何正拥'), ('雷有才', '雷友才'), ('王兴姣', '王兴蛟'), ('李金平', '金李平')]

# 收支项目 → 归类（用户原有的 19 个 + 新加的）
INC_ITEMS = [('工程款', '工程款收款'), ('工资退回', '工资发放'), ('退备用金', '备用金'), ('退保证金', '保证金押金'),
             ('银行借款', '银行借款'), ('利息收入', '其他收入'), ('其他收入', '其他收入'), ('股东投入', '股东投入'),
             ('个人借款-聂辉', '个人借款'), ('个人借款-王伟', '个人借款'), ('个人借款-张滨', '个人借款'), ('个人借款-张淑平', '个人借款'),
             ('往来款', '往来款'), ('其他应收款-代扣社保', '代收代付'), ('内部转账', '内部转账')]
EXP_ITEMS = [('项目-劳务费', '工资发放'), ('项目-材料费', '材料款'), ('项目-分包', '分包款'), ('项目-机械费', '机械费'),
             ('项目-期间费', '项目其他费用'), ('项目-保证金', '保证金押金'), ('办公室-办公费', '管理费用'), ('办公室-工资', '工资发放'),
             ('业务招待费', '管理费用'), ('税款', '税费'), ('农行-利息', '财务费用'), ('九江银行-利息', '财务费用'), ('手续费', '财务费用'),
             ('罚款', '营业外支出'), ('农行-本金', '银行借款'), ('个人借款-聂辉', '个人借款'), ('个人借款-王伟', '个人借款'),
             ('个人借款-张滨', '个人借款'), ('个人借款-张淑平', '个人借款'), ('往来款', '往来款'), ('其他应收款-代扣社保', '代收代付'),
             ('报销还款', '报销还款'), ('备用金', '备用金'), ('固定资产', '固定资产购置'), ('内部转账', '内部转账')]
ITEM_CAT = {}
for _n, _c in INC_ITEMS + EXP_ITEMS:
    ITEM_CAT.setdefault(_n, _c)
CAT_GS = {c[0]: c[3] for c in L.CATS}
_bad = sorted({c for c in ITEM_CAT.values() if c not in L.CAT_NAMES})
if _bad:
    raise RuntimeError(f'归类不在 layout.CATS 里：{_bad}')
SUPPLIER_CATS = ('材料款', '分包款', '机械费')

# 账户（原登记表的写法）
ACCOUNTS = [  # 名称, 类型, 所属人, 备注
    ('九江银行', '银行', '', '对公基本户'),
    ('中国农业银行', '银行', '', '贷款户（2026-03-10 贷款到账 361,000）'),
    ('中国工商银行', '银行', '', ''),
    ('建设银行', '银行', '', '1007 版【基础资料】新加的账户，原登记表还没有流水'),
    ('现金', '现金', '', ''),
    ('农民工专户', '专户', '', '甲方代发工资专户：收工程款、付工人工资，一收一付，余额应为 0'),
    ('九安账户', '专户', '', '九安公司代收代付户（2026-02-14 收 670,000、付筑强材料 420,000 和工人工资 250,000），余额应为 0'),
    ('温州爱思强贸易有限公司', '专户', '', '第三方代付户：2025-12-30 替公司付筑强 280,000，2026-02-14 退回，余额 0。'
                                     '另有同名材料供应商「温州爱思强贸易」（博河科创中心应付 11,100）'),
    ('聂辉微信', '个人户', '聂辉', ''),
    ('聂辉建行6680', '个人户', '聂辉', '原表 2025-07-04 付鑫城 50,000（李美坚），筑强退款 80,000＋40,000 打进此卡'),
    ('王伟微信', '个人户', '王伟', ''),
    ('张滨微信', '个人户', '张滨', ''),
    ('张滨支付宝', '个人户', '张滨', ''),
    ('张滨建行卡', '个人户', '张滨', '张滨个人明细：2026-06~09 由此卡转农行公户还贷款利息 3,800（登记表在农行记「个人借款-张滨」）；本卡无其他流水'),
]
ACC_NORM = {'九江银行（公）': '九江银行', '工商银行': '中国工商银行', '农业银行': '中国农业银行', '九安': '九安账户',
            '王伟微信转': '王伟微信', '王伟微信转账': '王伟微信', '王伟垫付': '王伟微信', '张滨微信转': '张滨微信',
            '聂辉建行私户': '聂辉建行6680', '退聂辉': '聂辉建行6680', '退温州爱思强贸易有限公司': '温州爱思强贸易有限公司',
            '农民工专户-代发工资': '农民工专户', '张滨代付租赁费': '张滨微信', '张滨垫付': '张滨微信', '叶焰平垫付': ''}
PERSONAL_PAY = {'聂辉微信': '聂辉微信', '王伟微信': '王伟微信', '王伟微信转账': '王伟微信', '王伟垫付': '王伟微信',
                '张滨代付租赁费': '张滨微信', '张滨垫付': '张滨微信'}


# ═════════════════════════════ 读源文件 ═════════════════════════════
def _wb(fn, ro=False):
    return openpyxl.load_workbook(P(fn), data_only=True, read_only=ro)


def _hdrmap(ws, row, stop='备注2', maxc=40):
    """表头行 → {规范化表头: 列号}，只取到 stop 那一列（右边的透视汇总块不要）"""
    out = {}
    for c in range(1, maxc + 1):
        h = hnorm(ws.cell(row, c).value)
        if h and h not in out:
            out[h] = c
        if stop and h == stop:
            break
    return out


def load_register():
    wb = _wb(F_REG)
    ws = wb['收支登记表']
    hm = _hdrmap(ws, 7, stop='备注')
    need = ['日期', '类别', '账户', '收支项目', '业务摘要', '收入金额', '支出金额', '项目名称', '负责人', '客户名称', '供应商名称',
            '销售合同', '采购合同', '备注']
    miss = [h for h in need if h not in hm]
    if miss:
        raise RuntimeError(f'收支登记表 表头缺：{miss}')
    rows = []
    for r in range(8, ws.max_row + 1):
        g = lambda h: ws.cell(r, hm[h]).value  # noqa: E731
        if g('日期') is None and g('账户') is None and g('收入金额') is None and g('支出金额') is None:
            continue
        rows.append(dict(r=r, 日期=todate(g('日期')), 类别=s(g('类别')), 账户=s(g('账户')), 收支项目=s(g('收支项目')),
                         摘要=s(g('业务摘要')), 收入=num(g('收入金额')), 支出=num(g('支出金额')), 项目原=s(g('项目名称')),
                         负责人=s(g('负责人')), 客户原=s(g('客户名称')), 供应商原=s(g('供应商名称')), 销售合同=s(g('销售合同')),
                         采购合同=s(g('采购合同')), 备注原=s(g('备注'))))
    tot = dict(收入=num(ws.cell(6, hm['收入金额']).value), 支出=num(ws.cell(6, hm['支出金额']).value))
    return rows, tot


def load_transfers():
    """收支登记表的外部链接缓存里留着原系统「内部转账」表（C 日期、E 转入、F 转出、G 金额）"""
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    out = []
    with zipfile.ZipFile(P(F_REG)) as z:
        for n in sorted(z.namelist()):
            if not (n.startswith('xl/externalLinks/externalLink') and n.endswith('.xml')):
                continue
            root = ET.fromstring(z.read(n))
            names = [e.get('val') for e in root.iter(ns + 'sheetName')]
            if '内部转账' not in names:
                continue
            sid = str(names.index('内部转账'))
            cells = {}
            for sd in root.iter(ns + 'sheetData'):
                if sd.get('sheetId') != sid:
                    continue
                for c in sd.iter(ns + 'cell'):
                    v = c.find(ns + 'v')
                    m = re.match(r'([A-Z]+)(\d+)$', c.get('r'))
                    if v is not None and m:
                        cells[(int(m[2]), m[1])] = v.text
            for row in sorted({k[0] for k in cells}):
                d, to, fr, amt = cells.get((row, 'C')), cells.get((row, 'E')), cells.get((row, 'F')), num(cells.get((row, 'G')))
                if d and re.fullmatch(r'\d+(\.\d+)?', d) and to and fr and amt:
                    out.append(dict(日期=todate(float(d)), 转入=to.strip(), 转出=fr.strip(), 金额=amt, 行=row))
    return out


def load_ar_file():
    wb = _wb(F_AR)
    out = {}
    for ws in wb.worksheets:
        if ws.title in ('基础信息表', '总表') or ws.title.startswith('Sheet'):
            continue
        hm = _hdrmap(ws, 5, stop=None, maxc=22)
        col = lambda *names: next((hm[n] for n in names if n in hm), None)  # noqa: E731
        c = dict(合同时间=col('合同时间'), 保证金=col('保证金'), 合同=col('合同总款', '合同金额'), 签证=col('签证金额（约）'),
                 已收=col('已收金额'), 收款日期=col('收款日期'), 收款账户=col('收款账户'), 扣款=col('扣款'), 开票=col('开票金额'),
                 开票日期=col('开票日期'), 税率=col('税率'), 备注=col('备注'), 项目=col('项目名称'), 客户=col('客户名称'))
        rows = []
        for r in range(6, ws.max_row + 1):
            g = lambda k: ws.cell(r, c[k]).value if c[k] else None  # noqa: E731
            vals = {k: g(k) for k in c}
            if all(num(vals[k]) in (None, 0) for k in ('保证金', '合同', '签证', '已收', '扣款', '开票')):
                continue
            rows.append(dict(r=r, 合同时间=vals['合同时间'], 保证金=num(vals['保证金']), 合同=num(vals['合同']), 签证=num(vals['签证']),
                             已收=num(vals['已收']), 收款日期=todate(vals['收款日期']), 收款账户=s(vals['收款账户']),
                             扣款=num(vals['扣款']), 开票=num(vals['开票']), 开票日期=todate(vals['开票日期']),
                             税率=num(vals['税率']), 备注=s(vals['备注']), 项目=s(vals['项目']), 客户=s(vals['客户'])))
        out[ws.title] = rows
    base = {}
    ws = wb['基础信息表']
    for r in range(2, ws.max_row + 1):
        if s(ws.cell(r, 1).value):
            base[s(ws.cell(r, 1).value)] = s(ws.cell(r, 2).value)
    return out, base


def _load_ap_book(fn, skip=('总表',)):
    """应付分表：每张供应商表 → 行列表（应付、付款、开票三段各自独立，只是写在同一行）；另返回总表里每家的应付/已付"""
    wb = _wb(fn)
    sheets = collections.OrderedDict()
    for ws in wb.worksheets:
        if ws.title in skip:
            continue
        hm = _hdrmap(ws, 2)
        if '日期' not in hm or '应付金额' not in hm:
            continue
        rows = []
        for r in range(3, ws.max_row + 1):
            g = lambda h: ws.cell(r, hm[h]).value if h in hm else None  # noqa: E731
            raw = {h: g(h) for h in ('日期', '项目名称', '应付金额', '已付金额', '付款日期', '付款账户', '收款账户', '已开票金额',
                                     '开票日期', '税率', '备注2')}
            if all(v in (None, '') for v in raw.values()):
                continue
            rows.append(dict(r=r, 日期=todate(raw['日期']), 日期原=raw['日期'], 项目=s(raw['项目名称']), 应付=num(raw['应付金额']),
                             应付原=raw['应付金额'], 已付=num(raw['已付金额']), 已付原=raw['已付金额'], 付款日期=todate(raw['付款日期']),
                             付款日期原=raw['付款日期'], 付款账户=s(raw['付款账户']), 收款账户=s(raw['收款账户']),
                             已开票=num(raw['已开票金额']), 已开票原=raw['已开票金额'], 开票日期=todate(raw['开票日期']),
                             税率=num(raw['税率']), 备注=s(raw['备注2'])))
        if rows:
            sheets[ws.title] = rows
    tot = {}
    ws = wb['总表']
    hr = next(r for r in range(1, 10) if hnorm(ws.cell(r, 2).value) == '供应商名称')
    for r in range(hr + 1, ws.max_row + 1):
        nm = s(ws.cell(r, 2).value)
        if nm and nm in sheets:
            tot[nm] = (num(ws.cell(r, 4).value) or 0.0, num(ws.cell(r, 5).value) or 0.0)
        if any(s(ws.cell(r, c).value).startswith('合计') for c in (1, 2, 3)):
            tot['__合计__'] = (num(ws.cell(r, 4).value) or 0.0, num(ws.cell(r, 5).value) or 0.0)
    return sheets, tot, wb


def load_mat():
    sheets, tot, wb = _load_ap_book(F_MAT)
    # 筑强右侧块：「项目｜已开票｜金额」
    extra = []
    ws = wb['筑强']
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=12, max_col=ws.max_column):
        for c in row:
            if s(c.value) == '已开票':
                pj, amt = s(ws.cell(c.row, c.column - 1).value), num(ws.cell(c.row, c.column + 1).value)
                if pj and amt:
                    extra.append(dict(项目=pj, 已开票=amt, r=c.row))
    return sheets, tot, extra


def load_fb():
    sheets, tot, _wbk = _load_ap_book(F_FB)
    return sheets, tot


def load_jx():
    wb = _wb(F_JX)
    ws = wb['Sheet1']
    hm = _hdrmap(ws, 2)
    day = []
    for r in range(3, ws.max_row + 1):
        g = lambda h: ws.cell(r, hm[h]).value if h in hm else None  # noqa: E731
        if g('日期') is None and g('供应商名称') is None:
            continue
        day.append(dict(r=r, 日期=todate(g('日期')), 项目=s(g('项目名称')), 供应商=s(g('供应商名称')), 种类=s(g('机械种类')),
                        单价=num(g('单价')), 上午=s(g('上午时间')), 下午=s(g('下午时间')), 小时=num(g('工作时长（小时）')),
                        应付=num(g('应付金额（元）')), 已付=num(g('已付金额（元）')), 付款日期=todate(g('付款日期')),
                        付款账户=s(g('付款账户')), 已开票=num(g('已开票金额')), 税率=num(g('税率')), 备注=s(g('备注2'))))
    ws = wb['Sheet3']
    hm = _hdrmap(ws, 2)
    summ = []
    for r in range(3, ws.max_row + 1):
        g = lambda h: ws.cell(r, hm[h]).value if h in hm else None  # noqa: E731
        nm = s(g('挖机班组/（运输）名称'))
        if not nm:
            continue
        summ.append(dict(r=r, 日期=todate(g('日期')), 日期原=s(g('日期')), 名称=nm, 项目=s(g('项目名称')), 应付=num(g('应付金额')),
                         已付=num(g('已付金额')), 付款日期=todate(g('付款日期')), 付款账户=s(g('付款账户')), 收款账户=s(g('收款账户')),
                         已开票=num(g('已开票金额')), 开票日期=todate(g('开票日期')), 税率=num(g('税率')), 备注=s(g('备注2'))))
    return day, summ


def load_1007():
    wb = _wb(F_U1007, ro=True)

    def rows(name, r0, maxc):
        return [(i, list(v)) for i, v in enumerate(wb[name].iter_rows(min_row=r0, max_col=maxc, values_only=True), start=r0)]
    out = {}
    out['往来单位'] = [dict(r=i, 名称=s(v[1]), 类型=s(v[2]), 别名1=s(v[3]), 别名2=s(v[4]), 全称=s(v[5]), 证件=s(v[6]),
                        手机=s(v[7]), 账号=s(v[8]), 开户行=s(v[9]), 期初欠薪=num(v[12]), 备注=s(v[13]))
                   for i, v in rows('往来单位', 4, 14) if s(v[1])]
    out['工资标准'] = [dict(r=i, 姓名=s(v[1]), 生效=todate(v[2]), 日薪=num(v[3]), 月薪=num(v[4]), 说明=s(v[5]))
                   for i, v in rows('工资标准', 4, 6) if any(x not in (None, '') for x in v[1:5])]
    kq = []
    for i, v in rows('考勤工资', 5, 22):
        if not s(v[2]):
            continue
        pairs = [(s(v[3 + 2 * k]), num(v[4 + 2 * k])) for k in range(4)]
        kq.append(dict(r=i, 月份=todate(v[1]), 姓名=s(v[2]), pairs=pairs, 单价=num(v[11]), 补贴=num(v[12]), 其他=num(v[13]),
                       个税=num(v[14]), 社保=num(v[15]), 备注=s(v[16]), 应发缓存=num(v[21])))
    out['考勤'] = kq
    out['项目档案'] = [dict(r=i, 简称=s(v[1]), 全称=s(v[2]), 甲方=s(v[3]), 合同=num(v[4]), 税率=num(v[9]), 开工=todate(v[10]),
                        完工=todate(v[11]), 状态=s(v[12]), 负责人=s(v[13]), 备注=s(v[18]))
                   for i, v in rows('项目档案', 4, 20) if s(v[1])]
    out['账户'] = [dict(r=i, 名称=s(v[6]), 类型=s(v[7]), 所属=s(v[8]), 期初=num(v[10]))
                 for i, v in rows('基础资料', 5, 16) if s(v[6])]
    out['公司简称'] = s(wb['基础资料']['C5'].value)
    out['垫付'] = [dict(r=i, 日期=todate(v[1]), 人=s(v[2]), 账户=s(v[3]), 摘要=s(v[4]), 项目=s(v[5]), 金额=num(v[7]))
                 for i, v in rows('垫付报销台账', 6, 8) if todate(v[1]) and num(v[7])]
    wb.close()
    return out


def load_nh():
    wb = _wb(F_NH)
    ws = wb['聂辉报销明细']
    hm = _hdrmap(ws, 2, stop='未报销金额')
    out = []
    for r in range(3, ws.max_row + 1):
        g = lambda h: ws.cell(r, hm[h]).value if h in hm else None  # noqa: E731
        d, amt = todate(g('日期')), num(g('支出金额'))
        if not d or not amt:
            continue
        out.append(dict(r=r, 日期=d, 账户=s(g('账户')), 摘要=s(g('业务摘要')), 金额=amt, 项目=s(g('项目名称')),
                        已开票=num(g('已开票金额')), 开票日期=todate(g('开票日期')), 开票单位=s(g('开票单位'))))
    return out


def load_personal():
    """个人借款及个人收支明细：只取各人「汇总/余款」用于核对"""
    wb = _wb(F_PERS)
    res = {}

    def find(sheet, label, col):
        ws = wb[sheet]
        for r in range(1, ws.max_row + 1):
            if s(ws.cell(r, 1).value) == label:
                return num(ws.cell(r, col).value), s(ws.cell(r, col + 1).value)
        return None, ''
    res['聂辉明细_个人借款余额'] = find('聂辉个人明细', '汇总', 4)[0]
    res['聂辉借款表_余款'] = find('聂辉个人借款', '余款', 2)
    res['王伟明细_个人借款余额'] = find('王伟个人明细', '汇总', 4)[0]
    res['王伟借款表_余款'] = find('王伟个人借款', '余款', 2)[0]
    res['张淑平_余款'] = find('张淑平个人借款', '余款', 2)[0]
    res['张滨明细_代付款'] = find('张滨个人明细', '汇总', 6)[0]
    return res


def load_arrears():
    wb = _wb(F_ARREAR)
    ws = wb['Sheet1']
    hm = _hdrmap(ws, 2, stop='剩余工资')
    out = []
    for r in range(3, ws.max_row + 1):
        nm, amt = s(ws.cell(r, hm['姓名']).value), num(ws.cell(r, hm['未发清金额']).value)
        if not nm or not amt:
            continue
        paid = []
        for h, c in hm.items():
            m = re.match(r'(\d{4})年(\d{1,2})月补发', h)
            v = num(ws.cell(r, c).value)
            if m and v:
                paid.append((f'{m[1]}-{int(m[2]):02d}', v))
        out.append(dict(姓名=nm, 未发清金额=amt, 补发=paid))
    return out


# ═════════════════════════════ 构建 ═════════════════════════════
class Ctx:
    def __init__(self):
        self.log, self.questions = [], []

    def L(self, msg):
        self.log.append(msg)

    def Q(self, msg):
        self.questions.append(msg)


def pj_canon(name, k):
    n = s(name)
    if not n:
        return ''
    if n in PSEUDO_PJ:
        return ''
    if n == '备用':
        return ''
    return PJ_ALIAS.get(n, n)


def cus_canon(name):
    n = s(name)
    return CUS_ALIAS.get(n, n)


def build_ctx(verbose=False):
    k = Ctx()
    reg, reg_tot = load_register()
    trans = load_transfers()
    arf, arbase = load_ar_file()
    mat, mat_tot, zq_inv = load_mat()
    fb, fb_tot = load_fb()
    jx_day, jx_sum = load_jx()
    u = load_1007()
    nh = load_nh()
    pers_led = load_personal()
    arrears_src = load_arrears()

    if len(reg) != 1062:
        k.L(f'⚠ 收支登记表读到 {len(reg)} 行（分析时是 1,062 行）')
    k.L(f'读入：收支登记表 {len(reg)} 行（第 {reg[0]["r"]}～{reg[-1]["r"]} 行）；外部「内部转账」{len(trans)} 笔；'
        f'1007 版考勤 {len(u["考勤"])} 行、工资标准 {len(u["工资标准"])} 行、往来单位 {len(u["往来单位"])} 个；聂辉报销明细 {len(nh)} 行')

    # ─────────── 人员 ───────────
    persons, per_idx = build_persons(k, u, reg)
    typo = {a: b for a, b in NAME_TYPO}
    names_sorted = sorted(set(per_idx) | set(typo), key=len, reverse=True)

    def find_name(text):
        for nm in names_sorted:
            if nm in text:
                return typo.get(nm, nm)
        return ''

    # ─────────── 项目、甲方 ───────────
    projects = build_projects(k, u, arf, reg)
    pj_names = {p['名称'] for p in projects}
    customers = [dict(名称=n, 简称=a, 税号='', 联系人='', 电话='', 开户='', 备注='') for n, a in CUS_LIST]
    cus_names = {c['名称'] for c in customers}
    for c in customers:
        if c['名称'] == '九江浔辉建筑工程有限公司':
            c['备注'] = '应收表德安养老中心页写成「九江浔辉建筑工程有限公司␠」（末尾多空格），已去掉'
        if c['名称'] == '共青农村商业银行股份有限公司':
            c['备注'] = '德安农商行总部、隆世光伏、共青办公楼的甲方。应收表写「共青农商行/共青农商银行」，1007 项目档案写「德安农商银行」；全称以营业执照为准'
        if c['名称'] == '德安县水利水电建筑工程公司':
            c['备注'] = '应收表、1007 往来单位写「…工程公司」，登记表吴山政府项目写「…工程有限公司」，暂并成一家，请核对全称'
        if c['名称'] == '九江欣埠置业有限公司':
            c['备注'] = '只有 2024 年投标保证金往来（付 100 万、2024-08-19 退回）'

    # ─────────── 现金流水：登记表 1062 行 ───────────
    cash = []
    for i, x in enumerate(reg):
        row = dict(日期=x['日期'], 类别=x['类别'], 账户=x['账户'], 收支项目=x['收支项目'], 摘要=x['摘要'],
                   收入=r2(x['收入']) if x['收入'] is not None else None, 支出=r2(x['支出']) if x['支出'] is not None else None,
                   项目='', 负责人=x['负责人'], 客户='', 供应商='', 人员='', 对方账户='', 费用归属='', 已开票=None, 开票日期=None, 备注='')
        notes = []
        if x['销售合同']:
            sc = x['销售合同']
            notes.append(f'销售合同 {sc}' if re.fullmatch(r'\d+(\.0)?', sc) else f'销售合同栏：{sc}')
        if x['采购合同']:
            notes.append(f'采购合同栏：{x["采购合同"]}')
        if x['备注原']:
            notes.append(x['备注原'])
        # 项目
        p0 = x['项目原']
        pj = pj_canon(p0, k)
        if p0 == '隆世光伏/共青办公楼/德祥工地' and x['收支项目'] == '工程款' and abs((x['收入'] or 0) - 85738) < 0.005:
            pj = '共青办公楼'
            notes.append('原项目「隆世光伏/共青办公楼/德祥工地」；按应收表这笔 85,738 是共青办公楼')
        if p0 in PSEUDO_PJ:
            notes.append(f'原项目栏「{p0}」')
        if p0 and pj and pj not in pj_names:
            raise RuntimeError(f'项目没有建档：{p0} → {pj}')
        row['项目'] = pj
        # 客户
        c0 = x['客户原']
        if c0:
            cc = cus_canon(c0)
            if cc in cus_names:
                row['客户'] = cc
            else:
                notes.append(f'原客户栏：{c0}')
        row['_notes'] = notes
        row['_i'] = i
        row['_src'] = ('登记表', x['r'])
        cash.append(row)
    # 第 9 行：收支项目、摘要、项目都空
    for row, x in zip(cash, reg):
        if not x['收支项目']:
            row['收支项目'] = '办公室-办公费'
            row['项目'] = COMPANY_PJ
            row['_notes'].append(f'原表收支项目、摘要、项目都空（收款方{x["供应商原"]}），暂按办公室-办公费')
            k.L(f'登记表第 {x["r"]} 行 {x["日期"]} {money(x["支出"])}：收支项目空 → 暂按「办公室-办公费」、项目「办公室」')
            k.Q(f'登记表第 {x["r"]} 行（{x["日期"]} 九江银行付 {x["供应商原"]} {money(x["支出"])}，原表收支项目和摘要都空），'
                f'暂记「办公室-办公费」，是培训费吗？')

    # ─────────── 供应商：把应付分表的付款配到登记表行 ───────────
    fpays = file_payments(k, mat, fb, jx_day, jx_sum)
    match_payments(k, reg, cash, fpays)

    # ─────────── 人员、供应商栏（其余） ───────────
    for row, x in zip(cash, reg):
        item, memo, p0 = x['收支项目'], x['摘要'], x['供应商原']
        person = ''
        if item.startswith('个人借款-'):
            person = item.split('-', 1)[1]
        elif item == '其他应收款-代扣社保' or x['项目原'] == '姚俊强':
            person = '姚俊强'
        elif item in ('项目-劳务费', '办公室-工资'):
            if row.get('_sup_match') and row['_sup_kind'] == '分包':
                person = ''
                row['_notes'].append(f'分包表把这笔记为「{row["_sup"]}」的分包付款（登记表记劳务费），人员不填，免得冲他考勤工资')
            elif not any(t in memo for t in NON_PERSON_MEMO):
                person = find_name(memo) or (find_name(p0) if p0 in per_idx or p0 in typo else '')
                if not person:
                    person = new_name(memo)
                    if person:
                        raise RuntimeError(f'工资行人名没进人员表：{x["r"]} {memo}')
        elif '报销' in memo or '备用金' in memo:
            person = find_name(memo)
        row['人员'] = person
        if person and person not in per_idx:
            raise RuntimeError(f'人员没建档：{person}（第 {x["r"]} 行）')
        # 供应商栏
        if row.get('_sup'):
            row['供应商'] = row['_sup']
            if p0 and p0 != row['_sup'] and SAME.get(p0) != row['_sup']:
                row['_notes'].append(f'实际收款：{p0}')
            continue
        sup = ''
        if p0:
            if p0 in SAME:
                sup = SAME[p0]
            elif p0 in per_idx or p0 in typo or p0 in ('张滨转入',):
                sup = ''
                if p0 not in (person,) and p0 not in ('张滨转入',):
                    row['_notes'].append(f'原供应商栏：{p0}')
            elif pj_canon(p0, k) in pj_names or p0 == '抚州资溪养老基地1号缕':
                sup = ''
                row['_notes'].append(f'原供应商栏填的是项目名「{p0}」')
            else:
                sup = p0
        if not sup and item in ('项目-分包', '项目-材料费', '项目-机械费', '项目-劳务费'):
            for nm, kws in list(SUB_KW.items()) + list(SUB_EXTRA.items()):
                if item == '项目-分包' and any(w in memo for w in kws):
                    sup = nm
                    break
            if not sup:
                for w, nm in MEMO_SUP.items():
                    if w in memo:
                        sup = nm
        row['供应商'] = sup

    # ─────────── 内部转账 10 笔 ───────────
    for t in trans:
        cash.append(dict(日期=t['日期'], 类别='支出', 账户=t['转出'], 收支项目='内部转账', 摘要=f'转入{t["转入"]}', 收入=None,
                         支出=r2(t['金额']), 项目='', 负责人='', 客户='', 供应商='', 人员='', 对方账户=t['转入'], 费用归属='',
                         已开票=None, 开票日期=None, 备注='', _notes=['来源：原系统「内部转账」表（收支登记表外部链接缓存）'],
                         _src=('内部转账', t['行'])))
    k.L(f'内部转账 {len(trans)} 笔，合计 {money(sum(t["金额"] for t in trans))}：每笔一行，账户＝转出、对方账户＝转入')

    # ─────────── 个人账户付、登记表没有的钱 ───────────
    before = acct_balances(cash)
    boss = boss_rows(k, nh, u, fpays, find_name, projects)
    cash.extend(boss)
    after = acct_balances(cash)

    src_cnt = collections.Counter(r['_src'][0] for r in cash)
    # 排序：按日期，同一天登记表在前（保持原顺序），再内部转账，再补录
    grp = {'登记表': 0, '内部转账': 1}
    cash.sort(key=lambda r: (r['日期'], grp.get(r['_src'][0], 2), r['_src'][1] if isinstance(r['_src'][1], int) else 0))
    for r in cash:
        r['备注'] = joinn(r.get('备注', ''), *r.pop('_notes', []))
        for kk in [kk for kk in r if kk.startswith('_')]:
            r.pop(kk)

    # ─────────── 供应商表 ───────────
    suppliers = build_suppliers(k, mat, fb, cash)

    # ─────────── 应收、应付、考勤、工资标准、欠薪 ───────────
    ar = build_ar(k, arf)
    ap = build_ap(k, mat, fb, jx_day, jx_sum, zq_inv)
    att = build_att(k, u, per_idx, pj_names)
    rates = build_rates(k, u, per_idx)
    # 截至 2026-02-28：用户表的补发从「2026年3月」起，1、2 月发的是当时的工资和补以前的欠薪，已经算在未发清里
    arrears = [dict(姓名=a['姓名'], 未发清金额=r2(a['未发清金额']), 截至=dt.date(2026, 2, 28),
                    备注=joinn('2025未发清工资.xlsx（补发从 2026 年 3 月起）', '已补发 ' + '、'.join(f'{m} {money(v)}' for m, v in a['补发']) if a['补发'] else ''))
               for a in arrears_src]
    for a in arrears_src:
        k.L(f'欠薪补发：{a["姓名"]} 未发清 {money(a["未发清金额"])}（截至 2026-02-28：原表补发从 2026 年 3 月起），表上已补发 ' +
            ('、'.join(f'{m} {money(v)}' for m, v in a['补发']) or '无'))

    # 名称一致性：录入表用到的名字都要在基本信息里
    sup_names = {x['名称'] for x in suppliers}
    acc_names = {a[0] for a in ACCOUNTS}
    for r in cash:
        for key, pool in (('项目', pj_names), ('客户', cus_names), ('供应商', sup_names), ('人员', set(per_idx)), ('账户', acc_names)):
            if r[key] and r[key] not in pool:
                raise RuntimeError(f'收支登记 {key}「{r[key]}」不在基本信息里：{r}')
        if r['对方账户'] and r['对方账户'] not in acc_names:
            raise RuntimeError(f'对方账户不在账户表：{r}')
        if r['收支项目'] not in ITEM_CAT:
            raise RuntimeError(f'收支项目没有归类：{r["收支项目"]}')
    for r in ap:
        if r['供应商'] not in sup_names or (r['项目'] and r['项目'] not in pj_names):
            raise RuntimeError(f'应付登记名字不在基本信息里：{r}')
    for r in ar:
        if r['项目'] not in pj_names:
            raise RuntimeError(f'应收登记项目不在项目档案：{r}')
    for r in att:
        if r['姓名'] not in per_idx or any(p and p not in pj_names for p, _d in r['pairs']):
            raise RuntimeError(f'考勤名字不在基本信息里：{r}')
    for r in rates:
        if r['姓名'] not in per_idx:
            raise RuntimeError(f'工资标准姓名不在人员信息：{r}')

    ctx = dict(
        params=build_params(k, u),
        accounts=[dict(名称=n, 类型=t, 所属人=o, 期初=0, 备注=b) for n, t, o, b in ACCOUNTS],
        inc_items=[dict(名称=n, 归类=c) for n, c in INC_ITEMS],
        exp_items=[dict(名称=n, 归类=c, 归属=CAT_GS.get(c, '')) for n, c in EXP_ITEMS],
        projects=projects, customers=customers, suppliers=suppliers, persons=persons, rates=rates,
        opening=[], cash=cash, ar=ar, ap=ap, att=att, inv=[], arrears=arrears,
    )
    k.L('期初余额：建账日 2023-05-01 早于所有数据（最早一笔是聂辉报销 2023-05-04，登记表从 2023-12-11 起），账户期初都是 0，不导期初')
    k.L('发票登记：没有电子税务局导出的真实数据，留空（销项开票记在应收登记、进项记在应付登记/收支登记的已开票）')
    for x in u['账户']:
        if x['名称'] not in acc_names and x['名称'] not in ('农业银行', '工商银行'):
            k.L(f'1007 版账户「{x["名称"]}」（{x["类型"]}）没有导：' +
                ('个人户改成按原表的微信/银行卡分开记' if x['类型'] == '个人户' else '新账套没有「票据」账户类型，承兑汇票原表直接记九江银行'))
    ctx['checks'] = run_checks(k, ctx, reg, before, after, arf, mat_tot, fb_tot, jx_day, jx_sum, fpays, pers_led, u, rates)
    ctx['checks']['收支登记来源'] = dict(src_cnt)
    si, so = r2(sum(x['收入'] or 0 for x in reg)), r2(sum(x['支出'] or 0 for x in reg))
    ac = collections.Counter(x['账户'] for x in reg)
    ctx['checks']['登记表原样'] = dict(行数=len(reg), 收入合计=si, 收入_表头J6=r2(reg_tot['收入']), 支出合计=so, 支出_表头K6=r2(reg_tot['支出']),
                                  农民工专户行数=ac['农民工专户'], 九安账户行数=ac['九安账户'],
                                  结果='✓' if len(reg) == 1062 and abs(si - reg_tot['收入']) < 0.01 and abs(so - reg_tot['支出']) < 0.01 else '✗')
    if ctx['checks']['登记表原样']['结果'] != '✓':
        raise RuntimeError(f'登记表导入后合计变了：{ctx["checks"]["登记表原样"]}')
    make_questions(k, ctx)
    ctx['log'], ctx['questions'] = k.log, k.questions
    return ctx


# ───────────── 人员 ─────────────
def new_name(memo):
    """登记表工资行里 1007 没登记的人名：去掉前缀、括号、奖金/退回/工资等字样后剩 2～3 个汉字"""
    m = re.sub(r'[（(].*?[）)]', '', memo).replace('❌️', '').replace('❌', '')
    for w in ('泥水修补工资', '挖沟，破路', '奖金', '退回', '工资'):
        m = m.replace(w, '')
    m = re.sub(r'\d+月份?', '', m)
    segs = [re.sub(r'补贴\d*', '', t).strip() for t in re.split(r'[-—]', m)]
    cands = [t for t in segs if cjk(t) and t not in NAME_STOP]
    return cands[-1] if cands else ''


def build_persons(k, u, reg):
    kq_names = collections.Counter(x['姓名'] for x in u['考勤'])
    rate_names = {x['姓名'] for x in u['工资标准'] if x['姓名']}
    rates = {x['姓名']: x for x in u['工资标准'] if x['姓名']}
    # 错别字：以考勤/工资标准/往来单位里出现的为准
    official = set(kq_names) | rate_names | {x['名称'] for x in u['往来单位']}
    regtxt = collections.Counter()
    for x in reg:
        if x['收支项目'] in ('项目-劳务费', '办公室-工资'):
            regtxt[x['摘要']] += 1
    for a, b in NAME_TYPO:
        na = sum(v for t, v in regtxt.items() if a in t)
        nb = sum(v for t, v in regtxt.items() if b in t)
        if b not in official and a in official:
            raise RuntimeError(f'错别字方向不对：{a}/{b}')
        k.L(f'人名统一：{a} → {b}（{b} 在 1007 考勤 {kq_names.get(b, 0)} 行/工资标准{"有" if b in rate_names else "无"}；'
            f'登记表工资行「{a}」{na} 行、「{b}」{nb} 行）' + ('；李金平 2024-06-08 付 5,000 后退回，当天改付金李平 5,000' if a == '李金平' else ''))
    persons, idx = [], {}

    def add(nm, kind, team='', note='', **kw):
        if nm in idx:
            return idx[nm]
        p = dict(姓名=nm, 类别=kind, 工种=team, 计薪='日薪', 过账='否', 期初欠薪=None, 身份证='', 电话='', 银行卡='', 备注=note)
        p.update(kw)
        r = rates.get(nm)
        if r and r['月薪'] and not r['日薪']:
            p['计薪'] = '月薪'
        idx[nm] = p
        persons.append(p)
        return p
    boss = {'聂辉': '老板', '王伟': '老板'}
    crew = {}                                          # 分包班组成员：登记表摘要「龚文坚劳务费-XX」「朱总分包劳务费-XX」
    for x in reg:
        m = re.match(r'(龚文坚|朱总)\S*劳务费-([一-龥]{2,3})', x['摘要'])
        if m:
            crew.setdefault(m[2], '九江朱总' if m[1] == '朱总' else '龚文坚')
    rng = collections.defaultdict(list)                # 1007 往来单位里各班组成员所在的行
    for x in u['往来单位']:
        if x['名称'] in crew:
            rng[crew[x['名称']]].append(x['r'])
    for x in u['往来单位']:
        if x['类型'] == '分包' and x['证件'] and x['名称'] not in crew:
            for tm, rs in rng.items():
                if min(rs) <= x['r'] <= max(rs):
                    crew[x['名称']] = tm
    for x in u['往来单位']:
        nm, t = x['名称'], x['类型']
        if '（' in nm or (t == '分包' and not x['证件']):
            continue                                  # 分包单位（恩旗老叶、张小朋（分包）…）→ 供应商信息
        if t not in ('工人', '管理人员', '分包', ''):
            continue
        team = ''
        if t == '分包':
            kind = '班组长' if nm == '龚文坚' else '工人'
            team = '龚文坚' if nm == '龚文坚' else crew.get(nm, '')
            note = f'{team}分包班组（钱走分包款，不记工资）' if team else '1007 往来单位类型写「分包」'
            team = f'{team}班组' if team else ''
        elif t == '管理人员':
            kind, note = boss.get(nm, '管理人员'), ''
        elif t == '':
            kind, note = '临时工', '1007 往来单位没填类型'
        else:
            kind, note = '工人', ''
        note = joinn(note, f'1007 备注：{x["备注"]}' if x['备注'] else '')
        bank = joinn(x['账号'], f'（{x["开户行"]}）' if x['开户行'] else '').replace('；（', '（')
        p = add(nm, kind, team, note=note, 身份证=x['证件'], 电话=x['手机'], 银行卡=bank)
        if x['期初欠薪']:
            p['期初欠薪'] = x['期初欠薪']
    for nm in list(kq_names) + sorted(rate_names):
        if nm not in idx:
            add(nm, '工人', note='1007 考勤/工资标准里有、往来单位没登记')
    # 特殊人员
    idx['姚俊强'].update(过账='是', 备注=joinn('工资、社保只在公司账上过一下（代扣社保记「其他应收款-代扣社保」）', idx['姚俊强']['备注']))
    for nm in ('朱金里', '梅景枝', '徐文元', '张友论'):
        if nm in idx:
            idx[nm]['工种'] = '油漆工'
    if '洪杰' in idx:
        idx['洪杰']['工种'] = '预算员'
        idx['洪杰']['备注'] = joinn('登记表「洪杰预算工资」「项目预算费-洪杰」记在项目-期间费', idx['洪杰']['备注'])
    if '涮油漆' in idx:
        idx['涮油漆']['备注'] = joinn('1007 考勤里的名字（刷油漆的人？），请核对', idx['涮油漆']['备注'])
    for a, b in NAME_TYPO:
        if b in idx:
            idx[b]['备注'] = joinn(f'登记表写作「{a}」', idx[b]['备注'])
    # 登记表工资行里出现、1007 没登记的人
    typo = {a: b for a, b in NAME_TYPO}
    known = sorted(set(idx) | set(typo), key=len, reverse=True)
    newp = collections.OrderedDict()
    for x in reg:
        if x['收支项目'] not in ('项目-劳务费', '办公室-工资'):
            continue
        memo = x['摘要']
        if any(t in memo for t in NON_PERSON_MEMO):
            continue
        if any(nm in memo for nm in known):
            continue
        nm = new_name(memo)
        if nm:
            newp.setdefault(nm, []).append(x['r'])
        else:
            k.L(f'登记表第 {x["r"]} 行工资「{memo}」认不出人名，人员留空')
    for nm, rows_ in newp.items():
        kind = '管理人员' if nm == '朱博文' else '临时工'
        add(nm, kind, note=f'1007 没登记；登记表工资行出现（第 {"、".join(map(str, rows_[:6]))} 行），没有考勤')
    if newp:
        k.L(f'人员：登记表工资行里有 {len(newp)} 个 1007 没登记的人，已加进人员信息（类别临时工）：{"、".join(newp)}')
    if '张淑平' in idx:
        idx['张淑平']['备注'] = joinn('借给公司 90,000、已还 50,000（个人借款-张淑平）', idx['张淑平']['备注'])
    k.L(f'人员信息 {len(persons)} 人（老板 2、管理人员 {sum(p["类别"] == "管理人员" for p in persons)}、'
        f'工人 {sum(p["类别"] == "工人" for p in persons)}、班组长 {sum(p["类别"] == "班组长" for p in persons)}、'
        f'临时工 {sum(p["类别"] == "临时工" for p in persons)}）；聂辉、王伟设为老板，张滨、张淑平、朱博文为管理人员，姚俊强过账＝是')
    return persons, idx


# ───────────── 项目 ─────────────
def build_projects(k, u, arf, reg):
    pa = {x['简称']: x for x in u['项目档案']}
    order = [x['简称'] for x in u['项目档案']]
    extra = ['共青办公楼', '温州盛为', '孙总鞋厂', '温州洪殿项目E02地块', '九江欣埠置业', '丰林高新区标准厂房']
    regnames = collections.defaultdict(set)
    for x in reg:
        if x['项目原'] and x['项目原'] not in PSEUDO_PJ:
            regnames[pj_canon(x['项目原'], k)].add(x['项目原'])
    for nm in sorted(regnames):
        if nm not in order and nm not in extra and nm != COMPANY_PJ:
            extra.append(nm)
    # 应收表：合同、保证金、合同时间
    ar_sheet_pj = {sh: PJ_ALIAS.get(sh, sh) for sh in arf}
    ar_info = {}
    for sh, rows in arf.items():
        pj = ar_sheet_pj[sh]
        main = next((r for r in rows if r['合同']), None)
        ar_info[pj] = dict(合同=main['合同'] if main else None, 保证金=sum(r['保证金'] or 0 for r in rows) or None,
                           合同时间=todate(main['合同时间']) if main else None, 客户=next((r['客户'] for r in rows if r['客户']), ''))
    jiafang = {
        '德安养老中心': '九江浔辉建筑工程有限公司', '科创中心': '德安县市政工程公司', '九安架空层': '九江浔辉建筑工程有限公司',
        '数字产业园': '江西中寰建筑装饰有限公司', '邹桥派出所': '德安县水利水电建筑工程公司', '抚州1号楼': '江西揽之建设工程有限公司',
        '抚州8号楼': '江西揽之建设工程有限公司', '抚州2号楼': '江西揽之建设工程有限公司', '博河科创中心': '德安县博河产业控股集团有限公司',
        '德安农商行总部': '共青农村商业银行股份有限公司', '隆世光伏': '共青农村商业银行股份有限公司',
        '共青办公楼': '共青农村商业银行股份有限公司', '赣粤高速': '江西赣粤高速公路工程有限责任公司',
        '中国联通项目': '联通（江西）产业互联网有限公司', '吴山政府项目': '德安县水利水电建筑工程公司',
        '温州盛为': '温州盛为装饰工程有限公司', '九江欣埠置业': '九江欣埠置业有限公司', '丰林高新区标准厂房': '九江浔辉建筑工程有限公司',
    }
    projects = [dict(名称=COMPANY_PJ, 全称='公司本身（办公室开支、管理人员工资、公司管理考勤）', 甲方='', 合同=None, 保证金=None, 质保比例=None,
                     质保金额=None, 质保到期=None, 开工=None, 完工=None, 状态='', 类型='公司', 负责人='', 税率=None,
                     备注='登记表项目栏「办公室」、1007 考勤「公司管理」都记到这里')]
    for nm in order + extra:
        if nm == COMPANY_PJ:
            continue
        p = pa.get(nm, {})
        info = ar_info.get(nm, {})
        alias = sorted(a for a in regnames.get(nm, set()) if a != nm)
        if p.get('全称') and p['全称'] != nm:
            alias.append(p['全称'])
        if nm == '共青办公楼':
            alias = ['隆世光伏/共青办公楼/德祥工地（登记表合写）']
        if nm == '数字产业园':
            alias.append('数字经济产业园（应收表页名）')
        notes = []
        if info.get('合同时间'):
            if okdate(info['合同时间']):
                notes.append(f'合同时间 {info["合同时间"]}（应收表）')
            else:
                notes.append('应收表合同时间录成 1900-01-25，未导')
                k.L(f'应收表「{nm}」合同时间是 {info["合同时间"]}（把 25 录成了日期格式），没导，记在项目备注')
        if nm == '隆世光伏':
            notes.append('登记表「隆世光伏/共青办公楼/德祥工地」的支出都记在这里（德祥工地没有单独数据）；收款 85,738 按应收表归共青办公楼')
        if nm == '九江欣埠置业':
            notes.append('只有投标保证金往来，净 0')
        if nm == '丰林高新区标准厂房':
            notes.append('保证金 200,000（2026-09-01 付、09-21 退回），保证金那行客户栏填的九江浔辉')
        if nm == '九安架空层':
            notes.append('九安账户 2026-02-14 收的 670,000 记本项目；同日九安账户付出的筑强材料 420,000、工资 250,000 原表记德安养老中心（照原表）')
        合同 = info.get('合同') or (p.get('合同') if nm not in ar_info else None)
        if nm == '邹桥派出所' and p.get('合同') and info.get('合同') and abs(p['合同'] - info['合同']) > 0.005:
            k.L(f'邹桥派出所合同额按应收表 {money(info["合同"])}、保证金 {money(info["保证金"])}（1007 项目档案是 {money(p["合同"])}，把保证金算进了合同额）')
        rate = p.get('税率') if p.get('税率') and abs(p['税率'] - 0.09) > 1e-9 else None
        projects.append(dict(名称=nm, 全称='、'.join(dict.fromkeys(alias)), 甲方=jiafang.get(nm, ''), 合同=r2(合同) if 合同 else None,
                             保证金=r2(info['保证金']) if info.get('保证金') else None, 质保比例=None, 质保金额=None, 质保到期=None,
                             开工=p.get('开工'), 完工=p.get('完工'), 状态='', 类型='工程', 负责人=p.get('负责人', ''),
                             税率=rate, 备注=joinn(*notes)))
    k.L(f'项目档案 {len(projects)} 个（含类型「公司」的「办公室」）：' + '、'.join(p['名称'] for p in projects))
    k.L('项目名对照：' + '；'.join(f'{a}→{b}' for a, b in PJ_ALIAS.items() if a != b) +
        '；伪项目「个人借款」「贷款」→ 空，「姚俊强」→ 空（人员＝姚俊强）；浔阳楼应付「备用」→ 空')
    k.L('合同额取应收表第一行（主合同）；德安养老中心 538,742.56、抚州8号楼 383,800 记在应收登记「补充协议」；开工/完工日期源表没有，留空（＝第一笔业务日）；'
        '状态留空（1007 版全填「在建」是上一版脚本的默认值）；税率只填跟默认 9% 不同的（抚州8号楼 13%，应收表开的是 13% 和 3%）；'
        '1007 项目档案的备注是上一版脚本写的提示，没导')
    return projects


# ───────────── 应付分表付款 → 登记表 ─────────────
def file_payments(k, mat, fb, jx_day, jx_sum):
    out = []
    for kind, book in (('分包', fb), ('材料', mat)):
        for sup, rows in book.items():
            for x in rows:
                if x['已付'] is None:
                    if s(x['已付原']):
                        k.L(f'{kind}「{sup}」第 {x["r"]} 行已付金额栏是文字「{s(x["已付原"])}」，不算付款')
                    continue
                if x['已付'] == 0:
                    continue
                d = x['付款日期'] if okdate(x['付款日期']) else None
                if d is None and kind == '分包' and sup == '龚文坚' and okdate(x['日期']):
                    d = x['日期']                      # 龚文坚表付款日期全空，用 A 列日期
                if s(x['付款日期原']) and not okdate(x['付款日期']):
                    k.L(f'{kind}「{sup}」第 {x["r"]} 行付款日期「{s(x["付款日期原"])}」不是日期，按无日期处理')
                out.append(dict(kind=kind, sup=sup, r=x['r'], date=d, amt=r2(x['已付']), acct=x['付款账户'], payee=x['收款账户'],
                                proj=x['项目'], note=x['备注']))
    for x in jx_day:
        if x['已付']:
            out.append(dict(kind='机械', sup=x['供应商'], r=x['r'], date=x['付款日期'] if okdate(x['付款日期']) else None,
                            amt=r2(x['已付']), acct=x['付款账户'], payee='', proj=x['项目'], note='机械费 Sheet1'))
    for x in jx_sum:
        nm = JX_NAME.get(x['名称'], x['名称'])
        if nm in ('杨工轮挖', '徐工轮挖'):
            continue                                  # 汇总表这两行是 Sheet1 台班的合计，付款也在 Sheet1
        if x['已付']:
            out.append(dict(kind='机械', sup=nm, r=x['r'], date=x['付款日期'] if okdate(x['付款日期']) else None,
                            amt=r2(x['已付']), acct=x['付款账户'], payee=x['收款账户'], proj=x['项目'], note='机械费 Sheet3'))
    for p in out:
        p['acct_n'] = ACC_NORM.get(p['acct'], p['acct'])
        p['m'] = []                                    # 配上的登记表行
    return out


def _text(x):
    return '|'.join((x['摘要'], x['供应商原'], x['销售合同'], x['采购合同'], x['客户原']))


def _compat(sup, kind, x):
    memo = x['摘要']
    hit = [n for n, kws in SUB_KW.items() if any(w in memo for w in kws)]
    if hit:
        return sup in hit
    kws = {'材料': MAT_KW, '分包': SUB_KW, '机械': JX_KW}[kind].get(sup, [])
    t = _text(x)
    return any(w in t for w in kws)


def match_payments(k, reg, cash, fpays):
    out_amt = [r2((x['支出'] or 0) - (x['收入'] or 0)) if x['类别'] == '支出' else None for x in reg]
    owner = {}                                         # 登记表下标 → 付款记录

    def take(i, p, how):
        owner[i] = (p, how)
        p['m'].append(i)
    order = sorted(fpays, key=lambda p: ({'分包': 0, '机械': 1, '材料': 2}[p['kind']], p['sup'], p['r']))
    # ① 金额相同、日期差 ≤5 天（没有日期的只看金额）
    for p in order:
        cands = []
        for i, x in enumerate(reg):
            if i in owner or out_amt[i] is None or abs(out_amt[i] - p['amt']) > 0.005 or not _compat(p['sup'], p['kind'], x):
                continue
            if p['date']:
                dd = abs((x['日期'] - p['date']).days)
                if dd > 5:
                    continue
            else:
                dd = 0
            cands.append(((x['账户'] != p['acct_n']) if p['acct_n'] else 0, dd, i))
        if cands:
            cands.sort()
            take(cands[0][2], p, '金额日期' if p['date'] else '只看金额')
    # ② 一笔付款＝登记表同账户 2～3 笔之和（日期差 ≤3 天）
    for p in order:
        if p['m'] or not p['date']:
            continue
        pool = [i for i, x in enumerate(reg) if i not in owner and out_amt[i] and out_amt[i] > 0 and abs((x['日期'] - p['date']).days) <= 3
                and _compat(p['sup'], p['kind'], x)]
        done = False
        for n in (2, 3):
            for combo in itertools.combinations(pool, n):
                if len({reg[i]['账户'] for i in combo}) == 1 and abs(sum(out_amt[i] for i in combo) - p['amt']) < 0.005:
                    for i in combo:
                        take(i, p, f'{n}笔合计')
                    k.L(f'付款配对（合计）：{p["kind"]}「{p["sup"]}」第 {p["r"]} 行 {money(p["amt"])} ＝ 登记表第 '
                        f'{"＋".join(str(reg[i]["r"]) for i in combo)} 行')
                    done = True
                    break
            if done:
                break
    # ③ 登记表一笔＝几家分表付款之和（例：叶焰平报销款 1,364.5 ＝ 启航 834 ＋ 零星 88＋5.5＋437）
    left = [p for p in order if not p['m']]
    for i, x in enumerate(reg):
        if i in owner or not out_amt[i] or out_amt[i] <= 0:
            continue
        cs = [p for p in left if not p['m'] and _compat(p['sup'], p['kind'], x) and (not p['date'] or (p['date'] - x['日期']).days <= 5)]
        if len({p['sup'] for p in cs}) < 2:
            continue
        found = None
        for n in range(2, min(5, len(cs)) + 1):
            for combo in itertools.combinations(cs, n):
                if abs(sum(p['amt'] for p in combo) - out_amt[i]) < 0.005:
                    found = combo
                    break
            if found:
                break
        if found:
            share = collections.Counter()
            for p in found:
                share[p['sup']] += p['amt']
                p['m'].append(i)
            main = share.most_common(1)[0][0]
            owner[i] = (next(p for p in found if p['sup'] == main), '几家合计')
            parts = '、'.join(f'{sp} {money(v)}' for sp, v in share.items())
            cash[i]['_notes'].append(f'这笔含 {parts}（分表分开记），整笔记在「{main}」')
            k.L(f'付款配对（一笔含几家）：登记表第 {x["r"]} 行 {money(out_amt[i])}「{x["摘要"]}」＝{parts}；一行只能填一家，记「{main}」，'
                f'核对表里 {main} 多 {money(sum(v for sp, v in share.items() if sp != main))}、其他家少同样的数')
    # ④ 冲回：同一天同账户同摘要的 +X −X +X，整串跟着配上的那笔走（净额仍是 X）
    nprop = 0
    for i in list(owner):
        p, how = owner[i]
        x = reg[i]
        j0 = i
        while j0 - 1 >= 0 and _same_run(reg[j0 - 1], x, out_amt[j0 - 1], out_amt[i]):
            j0 -= 1
        j1 = i
        while j1 + 1 < len(reg) and _same_run(reg[j1 + 1], x, out_amt[j1 + 1], out_amt[i]):
            j1 += 1
        run = list(range(j0, j1 + 1))
        if len(run) > 1 and abs(sum(out_amt[j] for j in run) - out_amt[i]) < 0.005:
            for j in run:
                if j not in owner:
                    owner[j] = (p, '冲回')
                    nprop += 1
    for i, (p, how) in owner.items():
        cash[i]['_sup'] = p['sup']
        cash[i]['_sup_kind'] = p['kind']
        cash[i]['_sup_match'] = True
    # 汇总日志
    by = collections.defaultdict(lambda: [0, 0, 0.0])
    for p in fpays:
        b = by[(p['kind'], p['sup'])]
        b[0] += 1
        if p['m']:
            b[1] += 1
            b[2] += p['amt']
    for kind in ('材料', '分包', '机械'):
        n = sum(v[0] for (kd, _s), v in by.items() if kd == kind)
        m = sum(v[1] for (kd, _s), v in by.items() if kd == kind)
        a = sum(v[2] for (kd, _s), v in by.items() if kd == kind)
        k.L(f'{kind}分表付款 {n} 笔，在登记表配上 {m} 笔（{money(a)}）；配上的行「供应商名称」改成分表的名字，实际收款方写进备注')
    k.L(f'冲回（+X −X +X）跟着配上的那笔一起记同一供应商：{nprop} 行')
    for p in fpays:
        if not p['m']:
            p['unmatched'] = True


def _same_run(y, x, ya, xa):
    return (y['日期'] == x['日期'] and y['账户'] == x['账户'] and y['摘要'] == x['摘要'] and ya is not None and xa is not None
            and abs(abs(ya) - abs(xa)) < 0.005)


# ───────────── 个人账户补录 ─────────────
def nh_item(memo, proj):
    if '吴群' in memo or '分包商' in memo:
        return '项目-分包'
    if any(w in memo for w in ('买材料', '镀锌管', '电线', '油漆', '板牙', '压力表')):
        return '项目-材料费'
    if proj and proj != COMPANY_PJ:
        return '项目-期间费'
    if any(w in memo for w in ('香烟', '买烟', '买酒', '请客', '招待')) or (memo.startswith('餐费-') and '员工' not in memo):
        return '业务招待费'
    return '办公室-办公费'


def boss_rows(k, nh, u, fpays, find_name, projects):
    rows = []

    def mk(d, acct, item, memo, amt, pj, sup='', inv=None, invd=None, note='', src=('补录', 0)):
        return dict(日期=d, 类别='支出', 账户=acct, 收支项目=item, 摘要=memo, 收入=None, 支出=r2(amt), 项目=pj, 负责人='', 客户='',
                    供应商=sup, 人员='', 对方账户='', 费用归属='', 已开票=inv, 开票日期=invd, 备注='', _notes=[note] if note else [],
                    _src=src)
    # 分表里个人账户付、登记表没有的
    pers = [p for p in fpays if p.get('unmatched') and p['acct'] in PERSONAL_PAY]
    other = [p for p in fpays if p.get('unmatched') and p['acct'] not in PERSONAL_PAY]
    # ① 聂辉报销明细 53 行（聂辉微信的原始记录，全部导入）
    nh_rows = []
    generic = []
    for x in nh:
        pj = pj_canon(x['项目'], k) if x['项目'] else ''
        item = nh_item(x['摘要'], pj)
        note = f'来源：聂辉报销明细第 {x["r"]} 行' + (f'；开票单位：{x["开票单位"]}' if x['开票单位'] else '')
        r = mk(x['日期'], '聂辉微信', item, x['摘要'], x['金额'], pj, inv=r2(x['已开票']) if x['已开票'] else None,
               invd=x['开票日期'], note=note, src=('聂辉报销明细', x['r']))
        r['_nh'] = x
        nh_rows.append(r)
        if item == '项目-分包':
            generic.append(r)
    # 分包付款（聂辉微信）与聂辉报销明细「微信转吴群」「转分包商」去重
    cand_exact = collections.Counter()                 # 每家有几笔「同一天同金额」能配上（平手时优先这家）
    pairs = []
    for r in generic:
        memo = r['摘要']
        named = [sp for sp, kws in SUB_KW.items() if any(w in memo for w in kws)]
        for p in pers:
            if PERSONAL_PAY[p['acct']] != '聂辉微信' or abs(p['amt'] - r['支出']) > 0.005 or not p['date']:
                continue
            if named and p['sup'] not in named:
                continue
            dd = abs((p['date'] - r['日期']).days)
            if dd <= 40:
                pairs.append((0 if named else 1, dd, p, r))
                if dd == 0:
                    cand_exact[p['sup']] += 1
    used_p, used_r = set(), set()
    for named, dd, p, r in sorted(pairs, key=lambda t: (t[0], t[1], -cand_exact[t[2]['sup']], t[2]['r'])):
        if id(p) in used_p or id(r) in used_r:
            continue
        used_p.add(id(p))
        used_r.add(id(r))
        p['dup'] = '聂辉报销明细'
        r['供应商'] = p['sup']
        r['项目'] = r['项目'] or pj_canon(p['proj'], k)
        r['_notes'].append(f'＝分包表「{p["sup"]}」第 {p["r"]} 行（{p["date"]} {money(p["amt"])}，收款 {p["payee"]}）')
        k.L(f'去重：聂辉报销明细第 {r["_nh"]["r"]} 行 {r["日期"]}「{r["摘要"]}」{money(r["支出"])} ＝ 分包表「{p["sup"]}」第 {p["r"]} 行 '
            f'{p["date"]} {money(p["amt"])}（聂辉微信）→ 只记一次（用报销明细的日期，供应商＝{p["sup"]}）' + (f'，日期差 {dd} 天' if dd else ''))
    for r in generic:
        if not r['供应商']:
            k.L(f'⚠ 聂辉报销明细第 {r["_nh"]["r"]} 行「{r["摘要"]}」没配上分包表，供应商留空')
    rows.extend(nh_rows)
    # ② 分表里其余个人账户付款 → 补录到个人账户
    pay_rows = []
    for p in pers:
        if p.get('dup'):
            continue
        acct = PERSONAL_PAY[p['acct']]
        d = p['date']
        note = f'来源：{"应付账款_分包" if p["kind"] == "分包" else "机械费"}「{p["sup"]}」第 {p["r"]} 行（登记表没有）；收款 {p["payee"] or "—"}'
        if p['kind'] == '机械':
            item = '项目-机械费'
        else:
            item = '项目-分包'
        if not d:
            if p['sup'] == '杨工轮挖':
                d = dt.date(2026, 1, 10)
                note += '；原表没写付款日期（Sheet1 写在 2025-10-13 那行、Sheet3 期间「25/11-26/1」），暂按杨工最后一天台班 2026-01-10'
            elif p['sup'] == '九江朱总':
                d = dt.date(2026, 2, 15)
                note += '；原表日期录成 1900-01-25，暂按朱总最后一批付款日 2026-02-15'
            else:
                raise RuntimeError(f'个人付款没日期：{p}')
            k.L(f'补录日期：{p["kind"]}「{p["sup"]}」第 {p["r"]} 行 {money(p["amt"])}（{p["acct"]}）原表无日期 → 暂记 {d}')
        if p['sup'] == '九江朱总' and '租赁' in p['acct']:
            note += '；实际是张滨替公司付给筑强的租赁费，分包表记在朱总名下（朱总因此多付 400），先照分包表记朱总'
        memo = {'恩旗老叶': '恩旗老叶分包-劳务费', '龚文坚': '龚文坚劳务费', '九江朱总': '张滨代付筑强租赁费（分包表记在朱总）',
                '张小朋': '张小朋分包劳务费'}.get(p['sup'], f'{p["sup"]}{"机械费" if p["kind"] == "机械" else "分包款"}')
        if p['payee'] and p['payee'] not in memo and p['kind'] == '分包':
            memo += f'-{p["payee"]}'
        r = mk(d, acct, item, memo, p['amt'], pj_canon(p['proj'], k), sup=p['sup'], note=note, src=('补录', 1000 + len(pay_rows)))
        r['_pay'] = p
        pay_rows.append(r)
    rows.extend(pay_rows)
    k.L(f'补录分表里个人账户付、登记表没有的付款 {len(pay_rows)} 笔：' + '；'.join(
        f'{a} {n} 笔 {money(v)}' for a, (n, v) in sorted(_agg(pay_rows, lambda r: r['账户']).items())))
    for p in other:
        k.L(f'⚠ 分表付款没在登记表找到、也不是个人账户付的：{p["kind"]}「{p["sup"]}」第 {p["r"]} 行 {p["date"] or "无日期"} {money(p["amt"])}'
            f'（{p["acct"] or "账户空"}）→ 不导（见问题）')
    # ③ 1007 垫付报销台账 14 行：跟上面已有的（同账户、±3 天同额，或 ±7 天内 2～3 笔之和）去重
    acc_rows = {a: [r for r in rows if r['账户'] == a] for a in ('王伟微信', '聂辉微信')}
    used = set()
    added = 0
    for x in u['垫付']:
        acct = x['账户']
        pool = [r for r in acc_rows.get(acct, []) if id(r) not in used]
        hit = None
        for r in pool:
            if abs(r['支出'] - x['金额']) < 0.005 and abs((r['日期'] - x['日期']).days) <= 3:
                hit = [r]
                break
        if not hit:
            near = [r for r in pool if abs((r['日期'] - x['日期']).days) <= 7]
            for n in (2, 3):
                for combo in itertools.combinations(near, n):
                    if abs(sum(r['支出'] for r in combo) - x['金额']) < 0.005:
                        hit = list(combo)
                        break
                if hit:
                    break
        if hit:
            for r in hit:
                used.add(id(r))
            k.L(f'去重：1007 垫付台账第 {x["r"]} 行 {x["日期"]} {x["人"]}「{x["摘要"]}」{money(x["金额"])} ＝ ' +
                '＋'.join(f'{r["日期"]}「{r["摘要"]}」{money(r["支出"])}' for r in hit) + ' → 不重复导')
            continue
        pj = pj_canon(x['项目'], k) or COMPANY_PJ
        memo = x['摘要']
        if pj == COMPANY_PJ:
            item = '业务招待费' if any(w in memo for w in ('客人', '招待', '买烟', '香烟')) else '办公室-办公费'
        elif '分包' in memo:
            item = '项目-分包'
        elif '机械' in memo:
            item = '项目-机械费'
        else:
            item = '项目-期间费'
        note = f'来源：1007 版垫付报销台账第 {x["r"]} 行（登记表没有）'
        if '草皮' in memo:
            note += '；含草皮树苗（材料）和临时工工资，金额待拆'
        if '间接费用' in memo:
            note += '；用途待确认（是不是红包/好处费）'
        rows.append(mk(x['日期'], acct, item, memo, x['金额'], pj, note=note, src=('1007垫付', x['r'])))
        added += 1
    k.L(f'1007 垫付台账 {len(u["垫付"])} 行，去重后导入 {added} 行')
    tot = _agg(rows, lambda r: r['账户'])
    k.L('个人账户补录合计：' + '；'.join(f'{a} {n} 行 {money(v)}' for a, (n, v) in sorted(tot.items())))
    k.L('聂辉报销明细收支项目：对外请客餐费、香烟、酒、请客水果→业务招待费；员工工作餐、聚餐、节日福利、办公用品→办公室-办公费；'
        '工地上的运费、过路费、油费、医疗、福利、技术员餐→项目-期间费；镀锌管、电线、油漆、板牙、压力表→项目-材料费；转吴群/转分包商→项目-分包')
    for r in rows:
        r.pop('_nh', None)
        r.pop('_pay', None)
    return rows


def _agg(rows, key):
    out = {}
    for r in rows:
        n, v = out.get(key(r), (0, 0.0))
        out[key(r)] = (n + 1, v + (r['支出'] or 0) - (r['收入'] or 0))
    return out


def acct_balances(cash, upto=CUTOFF):
    bal = collections.defaultdict(float)
    for r in cash:
        if r['日期'] and r['日期'] <= upto:
            net = (r['收入'] or 0) - (r['支出'] or 0)
            bal[r['账户']] += net
            if r['对方账户']:
                bal[r['对方账户']] -= net
    return {a: r2(v) for a, v in bal.items()}


# ───────────── 供应商信息 ─────────────
def build_suppliers(k, mat, fb, cash):
    used = collections.Counter(r['供应商'] for r in cash if r['供应商'])
    payees = collections.defaultdict(collections.Counter)
    for r in cash:
        m = re.search(r'实际收款：([^；]+)', r['备注'])
        if r['供应商'] and m:
            payees[r['供应商']][m[1]] += 1
    full = {}
    for a, b in SAME.items():
        if b in mat or b in fb:
            full.setdefault(b, []).append(a)
    sups = []

    def add(nm, typ, inv_unit='', rate=None, kind='', recv='', note=''):
        if any(x['名称'] == nm for x in sups):
            return
        sups.append(dict(名称=nm, 类型=typ, 开票单位=inv_unit, 税号='', 税率=rate, 发票类型=kind, 联系人='', 电话='', 收款账户=recv, 备注=note))
    special = {
        '浔阳楼电缆': dict(rate=0.13, kind='专票', note='经销商：付款多打给厂家（日昌、鄂缆、坚宝、双帝）；20 张票都是 13%'),
        '启航五金': dict(rate=0.01, kind='普票', note='已开票 10,358 是 1% 普票；另 11,298 没写税率'),
        '鑫城五金庄总': dict(note='付款打给李美坚（聂辉建行卡）、江西安迪消防器材；应付含「开票税费分摊」16,924.30'),
        '筑强': dict(note='扎账 2026-06-05（221,581.94）；按项目已开票 92,634 没有日期和税率'),
        '温州爱思强贸易': dict(note='同名单位在登记表里还被当成代付账户（见账户表）'),
        '零星供应商': dict(note='零星收款方：众贝源、共安机电（修路灯）、立邦、春虹、安南五金、敏芬、伦川、孙法炳'),
        '瑞安中拓拱水': dict(note='表名笔误，全称瑞安市中拓供水设备有限公司'),
        '佩力通风': dict(note='登记表另有恩旗老叶分包付给佩力的 60,000、10,000（记恩旗老叶，不冲佩力）'),
    }
    for nm in mat:
        sp = special.get(nm, {})
        inv_unit = ''
        fl = [a for a in full.get(nm, []) if a != nm]
        if len(fl) == 1:
            inv_unit = fl[0]
        elif nm == '四川久远':
            inv_unit = '四川久远智能消防设备有限公司'
        recv = '、'.join(p for p, _n in payees[nm].most_common() if p not in fl)
        note = sp.get('note', '')
        if len(fl) > 1:
            note = joinn(f'登记表写法：{"、".join(fl)}', note)
        add(nm, '材料', inv_unit, sp.get('rate'), sp.get('kind', ''), recv, note)
    sub_info = {
        '恩旗老叶': ('九江恩旗暖通设备有限公司', '包工包料；材料款直接付给佩力、拓齐、吉易通'),
        '图揽成善': ('江西图揽金属科技有限公司、江西成善科技有限公司', '两家公司合一张表；26,014 是 13% 专票'),
        '共安门业': ('江西共安门业有限公司', ''), '九江朱总': ('', '清包工，按工人逐个付'), '龚文坚': ('', '清包工，按工人逐个付；2026-04/06 两批走农民工专户'),
        '徐怀生': ('', '按工人逐个付'), '吴群': ('', '2026-04/05 付给向勇（代发工资）'), '祝冬冬': ('', '付款打给拓齐、东端、熊丽燕；数字产业园付了 42,000 但没有应付'),
        '张小朋': ('', '个人分包；也在考勤里（2026-04）'), '程章赐': ('', '个人分包；也在考勤里'), '吴聪': ('', ''),
    }
    for nm in fb:
        unit, note = sub_info.get(nm, ('', ''))
        recv = '、'.join(p for p, _n in payees[nm].most_common() if p not in unit)
        add(nm, '分包', unit, 0.13 if nm == '图揽成善' else None, '专票' if nm == '图揽成善' else '', recv, note)
    add('严程', '分包', note='登记表 2025-07-17「严程装监控」6,300（孙总鞋厂）；分包表没有这家')
    add('刘剑', '分包', note='登记表 2026-02-15「刘剑分包劳务费-陈雪寅」6,000（九安架空层）；分包表没有这家')
    add('九江迈腾建筑工程有限公司', '分包', '九江迈腾建筑工程有限公司', note='赣粤高速的劳务（登记表记项目-劳务费）')
    add('杨工轮挖', '机械', '九江翌年行建筑工程劳务有限公司', 0.01, '普票', '',
        '轮挖斗 150 元/时、炮头 250 元/时。翌年行是谁的开票单位待确认：机械费 Sheet1 把 5,000 付款和 5,000 发票记在杨工，Sheet3 记在徐工')
    add('徐工轮挖', '机械', note='1007 往来单位把徐工全称写成九江翌年行（待确认）')
    add('德安县恒建工程部', '机械', '德安县恒建工程部', 0.01, '普票', '舒文龙', '2026-06-01 德安农商行总部挖机 800')
    add('航帆增辉', '机械', note='邹桥派出所运输 1,000（王伟垫付）')
    add('章德吉货运', '机械', note='邹桥派出所运输 280（王伟垫付）')
    for nm, typ, note in (('江西安邦设备租赁有限公司', '机械', '2025-01-23「赣粤材料费」200,000 记在项目-机械费（性质待确认）'),
                          ('湖南开工智合数字技术有限公司', '机械', '抚州1号楼机械租赁 3,100'),
                          ('九江正鑫建筑劳务有限公司', '机械', '吴山政府项目机械费'), ('德安县弘升建筑工程有限公司', '机械', '吴山政府项目机械费')):
        add(nm, typ, nm, note=note)
    # 登记表里其余真实收款方（现买现付的材料、服务、保险、税局…）
    mat_like = {'德安小熊平价钢材店': '2024 年隆世光伏「付熊总材料费」，跟「九江安洋钢材熊总」是不是同一个熊总待确认',
                '湖北鄂缆电缆有限公司': '2025-10/11 四笔 32,274 在浔阳楼的发票里、不在浔阳楼已付里，先按现买现付',
                '西湖区荣晟消防设备营销中心': '', '江西柯源建材有限公司': '赣粤高速材料（130,760 那笔原表记在劳务费）', '温州萧通物资有限公司': '',
                '江苏野春种业有限公司': '草种子', '江西诚斯通信科技有限公司': '中国联通项目', '周丽丽': '水压机', '江西银枫消防科技有限公司': '付了又退，净 0',
                '南昌市西湖区联信大市场敏芬电缆器材经营部': ''}
    for nm in sorted(used):
        if any(x['名称'] == nm for x in sups):
            continue
        if nm in mat_like:
            add(nm, '材料', nm if not cjk(nm) else '', note=mat_like[nm])
        else:
            add(nm, '其他', nm if len(nm) > 4 else '', note='登记表收款方' + ('（含「九江财小保企业服务服限公司」「…有限公司␠」两种写法）' if '财小保' in nm else
                                                          '（含「中国平安保险股份有限公司」退费 703.20）' if '平安' in nm else ''))
    missing = [n for n in used if not any(x['名称'] == n for x in sups)]
    if missing:
        raise RuntimeError(f'供应商没建档：{missing}')
    k.L(f'供应商信息 {len(sups)} 家：材料 {sum(x["类型"] == "材料" for x in sups)}、分包 {sum(x["类型"] == "分包" for x in sups)}、'
        f'机械 {sum(x["类型"] == "机械" for x in sups)}、其他 {sum(x["类型"] == "其他" for x in sups)}；名称用应付分表的表名')
    k.L('供应商写法统一：四川久远智能消防设备有限公司/有限责任公司→四川久远；九江财小保企业服务服限公司、「…有限公司␠」→九江财小保企业服务有限公司；'
        '中国平安保险股份有限公司→中国平安财产保险股份有限公司广东分公司；朱总→九江朱总；九江翌年行→杨工轮挖；舒文龙→德安县恒建工程部；'
        '其余同一家公司的全称见供应商信息「开票单位」')
    return sups


# ───────────── 应收登记 ─────────────
def build_ar(k, arf):
    out = []
    for sh, rows in arf.items():
        pj = PJ_ALIAS.get(sh, sh)
        first = True
        for x in rows:
            if x['合同']:
                if not first:
                    d = todate(x['合同时间'])
                    note = joinn(x['备注'], '应收表合同栏第二行')
                    if not okdate(d):
                        # 抚州8号楼第 2 行合同 383,800 没有日期：取同一段第一张 3% 发票的日期（劳务部分）
                        d3 = sorted(r['开票日期'] for r in rows if r['开票'] and r['税率'] == 0.03 and r['开票日期'])
                        d = d3[0] if d3 else None
                        note = joinn(note, f'原表没有日期，暂按第一张 3% 发票日期 {d}')
                        k.L(f'应收：{pj} 补充协议 {money(x["合同"])} 原表无日期 → 暂记 {d}（该项目第一张 3% 劳务发票的日期）')
                    out.append(dict(日期=d, 项目=pj, 客户='', 类型='补充协议', 金额=r2(x['合同']), 税率=None, 发票号='', 备注=note))
                first = False
            if x['签证']:
                out.append(dict(日期=todate(x['合同时间']), 项目=pj, 客户='', 类型='签证', 金额=r2(x['签证']), 税率=None, 发票号='', 备注=''))
            if x['扣款']:
                out.append(dict(日期=None, 项目=pj, 客户='', 类型='扣款', 金额=r2(x['扣款']), 税率=None, 发票号='', 备注=''))
            if x['开票']:
                d = x['开票日期']
                if not okdate(d):
                    k.L(f'⚠ 应收：{pj} 开票 {money(x["开票"])} 日期不对（{d}），留空')
                    d = None
                out.append(dict(日期=d, 项目=pj, 客户='', 类型='开票', 金额=r2(x['开票']), 税率=x['税率'], 发票号='', 备注='应收账款.xlsx'))
    out.sort(key=lambda r: (r['日期'] or dt.date(2100, 1, 1), r['项目']))
    n = sum(r['类型'] == '开票' for r in out)
    k.L(f'应收登记 {len(out)} 行：开票 {n} 张 {money(sum(r["金额"] for r in out if r["类型"] == "开票"))}，'
        f'补充协议 {sum(r["类型"] == "补充协议" for r in out)} 行；收款不导（收支登记表里已有，逐项目一致）')
    return out


# ───────────── 应付登记 ─────────────
def build_ap(k, mat, fb, jx_day, jx_sum, zq_inv):
    out = []

    def add(**kw):
        row = dict(日期=None, 项目='', 供应商='', 类型='', 摘要='', 数量=None, 单价=None, 应付=None, 已开票=None, 开票日期=None,
                   税率=None, 发票类型='', 备注='')
        row.update(kw)
        out.append(row)

    def ptype(rate):
        return '专票' if rate and rate >= 0.09 else ('普票' if rate else '')
    # 材料：705 条送货 + 发票
    fixes = 0
    for sup, rows in mat.items():
        last_date = None
        for x in rows:
            if x['应付'] is not None and x['应付'] != 0:
                d = x['日期']
                note = x['备注']
                if d and d > CUTOFF and d.year == 2026 and d.month == 12:
                    nd = d.replace(year=2025)
                    k.L(f'应付日期更正：材料「{sup}」第 {x["r"]} 行 {d} → {nd}（未来日期，前后都是 2025-12）')
                    note = joinn(note, f'原表日期 {d}，按 {nd} 导入')
                    d = nd
                    fixes += 1
                if not okdate(d):
                    k.L(f'⚠ 材料「{sup}」第 {x["r"]} 行应付 {money(x["应付"])} 没有日期，用上一行日期 {last_date}')
                    d = last_date
                last_date = d or last_date
                pj = pj_canon(x['项目'], k)
                if x['项目'] == '备用':
                    note = joinn(note, '原表项目填「备用」')
                add(日期=d, 项目=pj, 供应商=sup, 类型='材料', 摘要='送货' if x['应付'] > 0 else '退货/扣减', 应付=r2(x['应付']), 备注=note)
            if x['已开票']:
                d = x['开票日期']
                note = f'第 {x["r"]} 行发票（项目照同一行，仅供参考）'
                if not okdate(d):
                    # 启航 11,298：没有开票日期，按同额付款日
                    pay = next((y for y in rows if y['已付'] and abs(y['已付'] - x['已开票']) < 0.005 and okdate(y['付款日期'])), None)
                    d = pay['付款日期'] if pay else None
                    note = joinn(note, f'原表无开票日期，按同额付款日 {d}' if d else '原表无开票日期')
                    k.L(f'应付发票：材料「{sup}」第 {x["r"]} 行 {money(x["已开票"])} 无开票日期 → 日期栏按同额付款日 {d}')
                add(日期=d, 项目=pj_canon(x['项目'], k), 供应商=sup, 类型='材料', 摘要='进项发票', 已开票=r2(x['已开票']),
                    开票日期=x['开票日期'] if okdate(x['开票日期']) else None, 税率=x['税率'], 发票类型=ptype(x['税率']), 备注=note)
            elif s(x['已开票原']):
                k.L(f'材料「{sup}」第 {x["r"]} 行已开票栏是文字「{s(x["已开票原"])}」，不算发票')
    for z in zq_inv:
        pj = pj_canon(z['项目'], k)
        ds = [x['日期'] for x in mat['筑强'] if pj_canon(x['项目'], k) == pj and x['应付'] and okdate(x['日期'])]
        d = max(ds) if ds else None
        add(日期=d, 项目=pj, 供应商='筑强', 类型='材料', 摘要='进项发票（按项目合计）', 已开票=r2(z['已开票']),
            备注=f'筑强表右侧第 {z["r"]} 行按项目的已开票合计，没有开票日期和税率；日期暂按该项目最后一次筑强送货 {d}')
        k.L(f'应付发票：筑强 {pj} 已开票 {money(z["已开票"])}（右侧块，无日期/税率）→ 记一行，日期取该项目最后一次筑强送货 {d}')
    # 分包：102 行应付 + 发票
    for sup, rows in fb.items():
        for x in rows:
            if x['应付'] is not None and x['应付'] != 0:
                d = x['日期']
                note = x['备注']
                if sup == '张小朋':
                    note = joinn(note, '张小朋表 A 列日期整列错一行（A 列＝下一行的付款日），应付日期待确认')
                if not okdate(d):
                    raise RuntimeError(f'分包应付没日期：{sup} {x}')
                add(日期=d, 项目=pj_canon(x['项目'], k), 供应商=sup, 类型='分包', 摘要=x['备注'] or '分包结算', 应付=r2(x['应付']),
                    备注=note if note != x['备注'] else '')
            if x['已开票']:
                d = x['开票日期'] if okdate(x['开票日期']) else x['日期']
                add(日期=d, 项目=pj_canon(x['项目'], k), 供应商=sup, 类型='分包', 摘要='进项发票', 已开票=r2(x['已开票']),
                    开票日期=x['开票日期'] if okdate(x['开票日期']) else None, 税率=x['税率'], 发票类型=ptype(x['税率']),
                    备注=joinn(f'第 {x["r"]} 行发票（项目照同一行）', '' if okdate(x['开票日期']) else f'原表无开票日期，日期栏按该行日期 {d}',
                               x['备注'] if x['备注'] == '开票未付' else ''))
    # 机械：87 行台班 + 汇总表 2026 年 3 笔 + 发票
    for x in jx_day:
        d = x['日期']
        note = ''
        sup = x['供应商']
        if sup == '徐工轮挖' and d and d < dt.date(2025, 10, 1):
            nd = d.replace(year=d.year + 1)
            note = f'原表日期 {d}，按 {nd} 导入（徐工台班 2025-12-30 起）'
            k.L(f'应付日期更正：机械「{sup}」第 {x["r"]} 行 {d} → {nd}')
            d = nd
        hrs = x['小时'] if x['小时'] is not None else _hours(x['上午'], x['下午'])
        amt = x['应付'] if x['应付'] is not None else (hrs or 0) * (x['单价'] or 0)
        add(日期=d, 项目=pj_canon(x['项目'], k), 供应商=sup, 类型='机械', 摘要=f'{x["种类"]} {joinn(x["上午"], x["下午"]).replace("；", " ")}'.strip(),
            数量=hrs, 单价=x['单价'], 应付=r2(amt), 备注=note)
        if x['已开票']:
            add(日期=x['付款日期'], 项目=pj_canon(x['项目'], k), 供应商=sup, 类型='机械', 摘要='进项发票', 已开票=r2(x['已开票']),
                税率=x['税率'], 发票类型=ptype(x['税率']), 备注=f'机械费 Sheet1 第 {x["r"]} 行（原表无开票日期，日期栏按付款日；Sheet3 把这张票记在徐工）')
    for x in jx_sum:
        nm = JX_NAME.get(x['名称'], x['名称'])
        if nm in ('杨工轮挖', '徐工轮挖'):
            continue
        add(日期=x['日期'], 项目=pj_canon(x['项目'], k), 供应商=nm, 类型='机械', 摘要='机械/运输', 应付=r2(x['应付']), 备注='机械费 Sheet3')
        if x['已开票']:
            add(日期=x['开票日期'] or x['付款日期'] or x['日期'], 项目=pj_canon(x['项目'], k), 供应商=nm, 类型='机械', 摘要='进项发票',
                已开票=r2(x['已开票']), 开票日期=x['开票日期'], 税率=x['税率'], 发票类型=ptype(x['税率']),
                备注='机械费 Sheet3（原表无开票日期，日期栏按付款日）' if not x['开票日期'] else '机械费 Sheet3')
    k.L(f'应付登记 {len(out)} 行：材料送货 {sum(1 for r in out if r["类型"] == "材料" and r["应付"] is not None)}、'
        f'分包 {sum(1 for r in out if r["类型"] == "分包" and r["应付"] is not None)}、机械 {sum(1 for r in out if r["类型"] == "机械" and r["应付"] is not None)}、'
        f'只有发票的 {sum(1 for r in out if r["应付"] is None)}；付款不从分表导（在收支登记里）')
    return out


def _hours(a, b):
    tot = 0.0
    for t in (a, b):
        m = re.match(r'(\d{1,2}):(\d{2})-(\d{1,2}):(\d{2})', t or '')
        if m:
            tot += (int(m[3]) * 60 + int(m[4]) - int(m[1]) * 60 - int(m[2])) / 60
    return tot or None


# ───────────── 考勤、工资标准 ─────────────
def build_att(k, u, per_idx, pj_names):
    out = []
    moved = 0
    for x in u['考勤']:
        d = x['月份']
        if not d:
            raise RuntimeError(f'考勤月份认不出：{x}')
        m1 = d.replace(day=1)
        if d.day != 1:
            moved += 1
        pairs = []
        for pj, days in x['pairs']:
            if not pj and not days:
                continue
            p2 = PJ_ALIAS.get(pj, pj)
            if not p2:
                k.L(f'⚠ 1007 考勤第 {x["r"]} 行 {x["姓名"]} 有天数 {days} 没项目')
            pairs.append([p2, days])
        extra, ded = x['补贴'] or 0, 0.0
        if x['其他']:
            if x['其他'] > 0:
                extra += x['其他']
            else:
                ded += -x['其他']
        ded += (x['个税'] or 0) + (x['社保'] or 0)
        out.append(dict(月份=m1, 姓名=x['姓名'], 工种='', pairs=pairs, 单价=x['单价'], 补贴=extra or None, 扣款=ded or None, 备注=x['备注']))
    k.L(f'考勤 {len(out)} 行（1007 版原样，一行一人一月，不合并）；月份统一成当月 1 号（{moved} 行原来是按天拖出来的日期序号）；「公司管理」→「办公室」')
    return out


def build_rates(k, u, per_idx):
    out, skip = [], []
    for x in u['工资标准']:
        if not x['姓名']:
            if x['日薪'] or x['月薪']:
                k.L(f'⚠ 工资标准第 {x["r"]} 行没有姓名')
            continue
        if not x['日薪'] and not x['月薪']:
            skip.append(x['姓名'])
            continue
        out.append(dict(姓名=x['姓名'], 生效=x['生效'], 日单价=x['日薪'], 月薪=x['月薪'], 备注=x['说明']))
    out.sort(key=lambda r: (r['姓名'], r['生效'] or dt.date(2000, 1, 1)))
    k.L(f'工资标准 {len(out)} 行；日薪月薪都空的 {len(skip)} 人没导（龚文坚班组、朱总班组、洪杰、油漆工等）：{"、".join(skip)}')
    k.L('工资标准里王贤锦、成峰、何正新、吴长友、金李平说明写「月薪」，但数填在日薪栏（320、370…），照日薪导')
    return out


def build_params(k, u):
    pv = {key: default for key, _lbl, default, _f, _d in L.PARAMS}
    pv['建账日'] = '2023-05-01'
    pv['考勤起算'] = '2025-03-01'
    short = u.get('公司简称') or ''
    if short:
        pv['公司'] = short
        k.L(f'公司名称：源文件里只找到简称「{short}」（1007 基础资料 C5、个人明细「竣辉农行公户」），全称、税号都没有')
    return pv


# ═════════════════════════════ 要用户确认的事 ═════════════════════════════
def make_questions(k, ctx):
    ch = ctx['checks']
    pa = ch['个人账户9月30日余额']
    sup = {x['供应商']: x for x in ch['每家应付已付核对']}

    def bal(nm):
        x = sup[nm]
        return money(x['应付_导入'] - x['其中能冲应付'])
    boss = collections.defaultdict(lambda: [0, 0.0])
    for r in ctx['cash']:
        if '来源：' in r['备注'] and r['账户'] in ('聂辉微信', '王伟微信', '张滨微信') and r['收支项目'] != '内部转账':
            boss[r['账户']][0] += 1
            boss[r['账户']][1] += (r['支出'] or 0) - (r['收入'] or 0)
    qs = [
        '【工行 281,148】2026-01-12 中国工商银行「工程款-揽之贷款」281,148：按应收表当抚州1号楼第二笔工程款导入（抚州1号楼已收 562,296）。'
        '如果其实是公司借款，要改收支项目，抚州1号楼已收会少 281,148。',
        '【老板个人账户补录】登记表里没有、按其他表补进个人账户的钱：' +
        '、'.join(f'{a} {n} 行 {money(v)}' for a, (n, v) in sorted(boss.items())) +
        '（聂辉报销明细 53 行、分包表个人付款 26 笔、1007 垫付台账去重后 10 行）。补完后 9/30 余额：' +
        '、'.join(f'{a} {money(v["补录后"])}' for a, v in pa.items() if v['补录后'] or v['补录前']) +
        '（负数＝公司欠他、正数＝他手上有公司的钱）。这些钱后来有没有报销过？余额对不对？',
        '【九安账户】2026-02-14 九安账户收 670,000 记九安架空层工程款，同日付出的筑强材料 420,000 和 11 个工人工资 250,000 原表记德安养老中心，'
        '照原表导入。是九安替德养付钱（要在两个项目间调），还是项目填错了？',
        '【机械·翌年行】杨工 30,000（王伟垫付）原表没写日期，暂记 2026-01-10。2025-11-10 付 5,000、2026-07-31 付 8,000 都是付给九江翌年行，'
        '现在都记杨工（机械费 Sheet1 把 5,000 和 1% 发票写在杨工行），但 Sheet3 和 1007 往来单位把翌年行算成徐工：翌年行是谁的开票单位？8,000 机械费表没记，要不要补台班？',
        '【平安保险 3,662.50】分包表恩旗老叶第 39 行「代付保险费」2026-01-01 九江银行 3,662.50，登记表九江银行里没有，没导'
        '（补进九江银行 9/30 余额就对不上 2,729.48）。实际是哪个账户付的？',
        f'【分表没有的材料付款】四川久远 2024 年隆世光伏 5 笔 42,313、天闽消防 2026-09-02 1,521：按同一家冲应付，所以四川久远显示应付余额 {bal("四川久远")}、'
        f'天闽 {bal("天闽消防")}（负数＝多付），隆世光伏、德养少了这部分材料成本。要不要在应付登记补这几批送货？'
        '湖北鄂缆 2025-10/11 四笔 32,274（浔阳楼的发票里有、浔阳楼已付里没有）先按现买现付记在「湖北鄂缆电缆有限公司」：算不算浔阳楼的货？',
        '【1007 垫付台账去重】按重复没导：第 6 行老叶 3,000＝分包表恩旗第 35 行；第 8 行 1,280＝机械航帆 1,000＋章德吉 280；'
        '第 18 行 5/9 员工聚餐 1,200＝聂辉报销明细 5/4 590＋5/6 320＋5/9 290；第 19 行 6/19 端午节买烟 3,520＝报销明细同日买香烟 3,520。对吗？',
        '【垫付用途】王伟「支付养老中心间接费用」4,000 是什么（红包？进德养成本还是待摊）？「付农商行草皮，树苗费用及临时工工资」480 怎么拆？',
        '【报销款（票未到）、备用金】张滨 6 笔 44,850、王伟 3 笔 30,000（办公室-办公费）和张滨备用金 60,000（项目-期间费，抚州1号楼）照原表记费用，'
        '人员栏标了是谁。要不要改成挂在个人名下（收支项目「报销还款」「备用金」），等票到了再记费用？',
        '【还农行利息 3,800】农行「张滨个人借款（支付贷款利息）」照原表记张滨借款（公司欠张滨 3,800）；张滨个人明细写「此款聂辉已转张滨」，'
        '原表张滨微信 9/30 余额 3,200 里有聂辉转来的 2,700（6/20 500、7/19 1,100、8/19 1,100），9/19 那 1,100 聂辉→张滨没记。要不要改成聂辉的借款？',
        '【工人退回 86,200】2026-05-14、06-18 工人退给张滨微信的 86,200（劳务费负数）又转给聂辉微信：这笔钱在聂辉那是「公司的钱」（个人账户余额为正），还是算公司还了聂辉的借款？',
        '【购酒款 20,000】2026-05-08 付咸宁群泰（净 20,000）原表记在个人借款-聂辉（＝还聂辉），照原表；聂辉个人明细 699,818 没算这笔。是还聂辉还是招待费？',
        '【隆世光伏/共青办公楼/德祥工地】收款 85,738 按应收表归共青办公楼，其余收款 600,000 和全部支出（125 行）归隆世光伏。共青办公楼、德祥工地的成本要分出来吗？',
        '【浔阳楼「备用」4,230】2026-06-27 应付项目写「备用」，导入时项目空着：属于哪个项目？',
        '【日期更正】筑强第 114 行 2026-12-31 改成 2025-12-31；徐工 4 天台班 2025-01-02～05 改成 2026-01-02～05。对吗？',
        '【缺开票日期/税率的票】筑强按项目已开票 92,634（邹桥 23,555.71、数字产业园 33,897.01、科创 35,181.28）、启航 11,298、'
        '图揽成善 4 张（5,093、7,772、2,230、5,009）、机械 5,000、800：请补开票日期和税率。其他供应商有没有收票（外账会计的进项导出）？',
        '【抚州8号楼 383,800】应收表合同栏第二行 383,800 记成「补充协议」，日期暂按第一张 3% 发票 2026-04-29：是补充协议，还是另一份劳务合同（开 3% 票）？',
        '【张小朋】分包表 A 列日期整列错一行，应付 52,410 暂记 2026-01-25。2025-09-18 那 5,000 登记表记劳务费（德养）、分包表算分包款（数字产业园），'
        '供应商栏已填张小朋，但劳务费冲不了应付（张小朋显示还欠 7,150）。按哪个？张小朋、程章赐既做分包又在考勤里，是两份活吗？',
        '【祝冬冬】数字产业园付了 42,000（熊丽燕），可应付里数字产业园是 0：是漏登了结算，还是预付？',
        '【朱总 400】分包表朱总第 16 行「张滨代付租赁费」400（收款筑强，日期录成 1900-01-25），暂记张滨微信 2026-02-15、供应商朱总：应该挂朱总还是筑强？',
        '【小差额】龚文坚多付 3,845、佩力多付 1,673、共安门业差 0.53、恒讯差 0.05、筑强第 8 行 0.91 登记表里找不到：要不要平掉？',
        '【邓付贞】2025 未发清 19,430（用户表）vs 按考勤和流水倒算 15,048，差 4,382：以哪个为准？其他工人 2025 年有没有没发清的？',
        '【姚俊强】设成过账人员（考勤不算工资成本）。4,500/月的工资实际发没发？7～9 月社保（1,019.48×3）退了没有？现在他欠公司 3,056.88。',
        '【邹桥 4,554】项目档案「保证金」填了 4,554（合同 91,006.56），2025-11-10 付款记项目-保证金：是我方交的保证金，还是甲方扣的质保金？',
        '【项目资料】各项目开工、完工日期和状态（导入时都空着）；赣粤高速（已收 1,190,284.50）、中国联通项目（650,862.40）、吴山政府项目（113,800）、'
        '温州盛为、星辰宾馆、孙总鞋厂、恒科华府的合同额和开票。',
        '【公司】公司全称和税号（基本信息①只找到简称「竣辉」）。',
        '【甲方名称】共青农商行的全称？德安农商行总部的发票抬头「农村商业银行股份有限公司德安县总部」是不是另一家？'
        '「德安县水利水电建筑工程公司」和「…有限公司」哪个对？丰林高新区标准厂房的甲方是不是九江浔辉（保证金那行客户栏填的）？',
        '【标记】负责人栏和摘要里的「❌」「无考勤」是什么意思（没票？没考勤？），照原样留着。',
        '【聂辉报销明细】2026-06-03「共青厂房预算费」10,000 暂记办公室-办公费：是哪个项目的（丰林高新区标准厂房？）？'
        '对外请客的餐费、香烟、酒记成「业务招待费」、员工餐记「办公室-办公费」，这样分可以吗？',
        '【熊总】德安小熊平价钢材店（2024 年隆世光伏「付熊总材料费」194,436）跟「九江安洋钢材熊总」是不是同一个人？',
        '【分类可疑】2025-01-23「赣粤材料费」200,000（江西安邦设备租赁）记在项目-机械费、「柯源建材」130,760 和迈腾 4 笔记在项目-劳务费：照原表，分类对吗？',
        '【考勤】1007 考勤有 29 个人月天数超过当月天数（年后 2 月并进 3 月等），照原样导入了；要拆到 2 月吗？',
    ]
    k.questions[:0] = qs


# ═════════════════════════════ 核对 ═════════════════════════════
def run_checks(k, ctx, reg, before, after, arf, mat_tot, fb_tot, jx_day, jx_sum, fpays, pers_led, u, rates):
    ch = collections.OrderedDict()
    cash, ap, ar = ctx['cash'], ctx['ap'], ctx['ar']
    # 1 账户余额
    bal = acct_balances(cash)
    bank = {}
    fails = []
    for a, exp in BANK_EXPECT.items():
        got = bal.get(a, 0.0)
        ok = abs(got - exp) < 0.005
        bank[a] = dict(应为=exp, 算出=got, 结果='✓' if ok else '✗ FAIL')
        if not ok:
            fails.append(f'{a} 算出 {money(got)}，应为 {money(exp)}')
    ch['公司账户9月30日余额'] = bank
    personal = [a for a, t, _o, _b in ACCOUNTS if t in ('个人户',)] + ['温州爱思强贸易有限公司']
    ch['个人账户9月30日余额'] = {a: dict(补录前=before.get(a, 0.0), 补录后=after.get(a, 0.0)) for a in personal}
    ch['全部账户9月30日余额'] = {a: bal.get(a, 0.0) for a, *_r in ACCOUNTS}
    if fails:
        raise RuntimeError('✗✗✗ 银行/现金余额对不上：' + '；'.join(fails))
    # 2 应付：分表合计 vs 应付登记
    apsum = collections.defaultdict(float)
    apsup = collections.defaultdict(float)
    for r in ap:
        if r['应付'] is not None:
            apsum[r['类型']] += r['应付']
            apsup[r['供应商']] += r['应付']
    jx_total = sum(x['应付'] or 0 for x in jx_day) + sum(x['应付'] or 0 for x in jx_sum if JX_NAME.get(x['名称'], x['名称']) not in ('杨工轮挖', '徐工轮挖'))
    file_tot = {'材料': mat_tot['__合计__'][0], '分包': fb_tot['__合计__'][0], '机械': jx_total}
    ch['应付合计（应付登记 vs 分表）'] = {t: dict(分表=r2(file_tot[t]), 导入=r2(apsum[t]), 结果='✓' if abs(file_tot[t] - apsum[t]) < 0.01 else '✗')
                                for t in file_tot}
    for t, v in ch['应付合计（应付登记 vs 分表）'].items():
        if v['结果'] != '✓':
            raise RuntimeError(f'应付合计对不上：{t} {v}')
    # 3 每家已付：收支登记里供应商＝这家的付款 vs 分表已付
    file_paid = {}
    for nm, (a, b) in list(mat_tot.items()) + list(fb_tot.items()):
        if nm != '__合计__':
            file_paid[nm] = (a, b)
    jxp = collections.defaultdict(lambda: [0.0, 0.0])
    for x in jx_day:
        jxp[x['供应商']][0] += x['应付'] or 0
        jxp[x['供应商']][1] += x['已付'] or 0
    for x in jx_sum:
        nm = JX_NAME.get(x['名称'], x['名称'])
        if nm not in ('杨工轮挖', '徐工轮挖'):
            jxp[nm][0] += x['应付'] or 0
            jxp[nm][1] += x['已付'] or 0
    for nm, (a, b) in jxp.items():
        file_paid[nm] = (a, b)
    paid_all = collections.defaultdict(float)
    paid_ap = collections.defaultdict(float)
    for r in cash:
        if r['供应商']:
            net = (r['支出'] or 0) - (r['收入'] or 0)
            paid_all[r['供应商']] += net
            if ITEM_CAT.get(r['收支项目']) in SUPPLIER_CATS:
                paid_ap[r['供应商']] += net
    reasons = {
        '筑强': '分表第 8 行 0.91 登记表里找不到',
        '启航五金': '登记表 2026-02-15「叶焰平报销款」1,364.50 一笔含启航 834 和零星 530.50，整笔记在启航',
        '零星供应商': '零星的 88＋5.5＋437 记在启航那笔里；另「科创中心修路灯」800 登记表记项目-期间费，冲不了应付',
        '天闽消防': '登记表 2026-09-02「气体灭火装置材料款」1,521 分表没有（没记应付）',
        '四川久远': '登记表 2024 年隆世光伏付「四川久远智能消防设备有限责任公司」5 笔 42,313，分表没有这批送货',
        '恩旗老叶': '分表第 39 行 2026-01-01 九江银行「代付保险费」3,662.50 登记表九江银行里没有，没补（见问题）',
        '杨工轮挖': '登记表 2026-07-31 付翌年行「科创中心挖机费」8,000，机械费表没记',
        '张小朋': '2025-09-18 那 5,000 登记表记的是项目-劳务费（工资），冲不了应付',
    }
    rows = []
    for nm, (fa, fp) in file_paid.items():
        got, got_ap = r2(paid_all.get(nm, 0)), r2(paid_ap.get(nm, 0))
        diff = r2(got - fp)
        rows.append(dict(供应商=nm, 应付_分表=r2(fa), 应付_导入=r2(apsup.get(nm, 0)), 已付_分表=r2(fp), 已付_收支登记=got,
                         其中能冲应付=got_ap, 差额=diff, 说明=reasons.get(nm, '') if abs(diff) >= 0.005 or abs(got_ap - got) >= 0.005 else ''))
    ch['每家应付已付核对'] = rows
    # 4 项目：合同、补充、开票、收款
    pjrows = []
    recv = collections.defaultdict(float)
    for r in cash:
        if ITEM_CAT.get(r['收支项目']) == '工程款收款':
            recv[r['项目']] += (r['收入'] or 0) - (r['支出'] or 0)
    invsum = collections.defaultdict(float)
    supp = collections.defaultdict(float)
    for r in ar:
        if r['类型'] == '开票':
            invsum[r['项目']] += r['金额']
        elif r['类型'] == '补充协议':
            supp[r['项目']] += r['金额']
    pjmap = {p['名称']: p for p in ctx['projects']}
    for sh, xs in arf.items():
        pj = PJ_ALIAS.get(sh, sh)
        f_con = sum(x['合同'] or 0 for x in xs)
        f_inv = sum(x['开票'] or 0 for x in xs)
        f_rec = sum(x['已收'] or 0 for x in xs)
        con = (pjmap[pj]['合同'] or 0) + supp[pj]
        pjrows.append(dict(项目=pj, 合同加补充_应收表=r2(f_con), 合同加补充_导入=r2(con), 保证金=pjmap[pj]['保证金'],
                           开票_应收表=r2(f_inv), 开票_导入=r2(invsum[pj]), 已收_应收表=r2(f_rec), 工程款_收支登记=r2(recv[pj]),
                           结果='✓' if abs(f_con - con) < 0.01 and abs(f_inv - invsum[pj]) < 0.01 and abs(f_rec - recv[pj]) < 0.01 else '✗'))
    for pj in sorted(recv):
        if pj not in [x['项目'] for x in pjrows]:
            pjrows.append(dict(项目=pj or '（空）', 合同加补充_应收表=None, 合同加补充_导入=pjmap.get(pj, {}).get('合同'), 保证金=None,
                               开票_应收表=None, 开票_导入=r2(invsum[pj]), 已收_应收表=None, 工程款_收支登记=r2(recv[pj]), 结果='应收表没有这个项目'))
    ch['项目应收核对'] = pjrows
    for x in pjrows:
        if x['结果'] == '✗':
            raise RuntimeError(f'项目应收核对不过：{x}')
    # 5 个人借款
    loans = collections.defaultdict(float)
    for r in cash:
        if ITEM_CAT.get(r['收支项目']) == '个人借款':
            loans[r['人员']] += (r['收入'] or 0) - (r['支出'] or 0)
    nl = pers_led
    ledger = {'聂辉': (r2((nl['聂辉明细_个人借款余额'] or 0) - 20000), f'聂辉个人明细合计 {money(nl["聂辉明细_个人借款余额"])}，减购酒款 20,000（原表算还聂辉）'),
              '王伟': (nl['王伟明细_个人借款余额'], f'王伟个人明细合计（王伟个人借款表余款 {money(nl["王伟借款表_余款"])} 少记 2026-09-02 还 13,000）'),
              '张淑平': (nl['张淑平_余款'], '张淑平个人借款表余款'),
              '张滨': (nl['张滨明细_代付款'], '张滨个人明细「代付款」合计（代付农行利息）')}
    ch['个人借款（公司欠）'] = {p: dict(收支登记=r2(loans.get(p, 0)), 个人明细=ledger[p][0], 说明=ledger[p][1],
                                 结果='✓' if ledger[p][0] is not None and abs(loans.get(p, 0) - ledger[p][0]) < 0.01 else '✗') for p in ledger}
    # 6 考勤应发（按新规则复算 vs 1007 缓存）
    rate_by = collections.defaultdict(list)
    for r in rates:
        rate_by[r['姓名']].append(r)
    pm_days = collections.defaultdict(float)
    for r in ctx['att']:
        pm_days[(r['姓名'], r['月份'])] += sum(d or 0 for _p, d in r['pairs'])
    tot, tot_pj = 0.0, collections.defaultdict(float)
    for r in ctx['att']:
        me = (r['月份'].replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
        std = [x for x in rate_by.get(r['姓名'], []) if not x['生效'] or x['生效'] <= me]
        std = std[-1] if std else None
        days = sum(d or 0 for _p, d in r['pairs'])
        if r['单价']:
            amt = r['单价'] * days
        elif std and std['月薪'] and not std['日单价']:
            amt = std['月薪'] * (days / pm_days[(r['姓名'], r['月份'])] if pm_days[(r['姓名'], r['月份'])] else 1)
        elif std:
            amt = (std['日单价'] or 0) * days
        else:
            amt = 0
        amt += (r['补贴'] or 0) - (r['扣款'] or 0)
        tot += amt
    cache = sum(x['应发缓存'] or 0 for x in u['考勤'])
    ch['考勤应发'] = dict(按新规则复算=r2(tot), 原1007缓存=r2(cache), 结果='✓' if abs(tot - cache) < 0.01 else '✗（差额来自 1007 原公式口径）')
    # 7 条数
    ch['条数'] = {key: len(ctx[key]) for key in ('accounts', 'inc_items', 'exp_items', 'projects', 'customers', 'suppliers', 'persons',
                                                'rates', 'opening', 'cash', 'ar', 'ap', 'att', 'inv', 'arrears')}
    k.L('核对：九江银行 {0}、农行 {1}、工行 {2}、现金 {3}、农民工专户 {4}、九安账户 {5}（9/30）'.format(
        *(money(bal.get(a, 0)) for a in BANK_EXPECT)))
    return ch


# ═════════════════════════════ 运行 ═════════════════════════════
def _json(o):
    if isinstance(o, (dt.date, dt.datetime)):
        return o.isoformat()[:10]
    raise TypeError(type(o))


def main():
    ctx = build_ctx()
    with open(os.path.join(HERE, '_ctx.json'), 'w', encoding='utf-8') as f:
        json.dump(ctx, f, ensure_ascii=False, indent=1, default=_json)
    ch = ctx['checks']
    print('== 条数 ==')
    print('  ' + '，'.join(f'{a} {b}' for a, b in ch['条数'].items()))
    print('  收支登记来源：' + '，'.join(f'{a} {b}' for a, b in ch['收支登记来源'].items()))
    print('== 公司账户 9/30 余额 ==')
    for a, v in ch['公司账户9月30日余额'].items():
        print(f'  {v["结果"]} {a}: {money(v["算出"])}（应为 {money(v["应为"])}）')
    print('== 个人账户 9/30 余额（补录前 → 补录后） ==')
    for a, v in ch['个人账户9月30日余额'].items():
        print(f'  {a}: {money(v["补录前"])} → {money(v["补录后"])}')
    print('== 应付合计 ==')
    for t, v in ch['应付合计（应付登记 vs 分表）'].items():
        print(f'  {v["结果"]} {t}: 分表 {money(v["分表"])}，导入 {money(v["导入"])}')
    print('== 每家已付（收支登记 vs 分表，只列有差的） ==')
    for x in ch['每家应付已付核对']:
        if abs(x['差额']) >= 0.005 or abs(x['其中能冲应付'] - x['已付_收支登记']) >= 0.005:
            print(f'  {x["供应商"]}: 分表已付 {money(x["已付_分表"])}，收支登记 {money(x["已付_收支登记"])}（能冲应付 {money(x["其中能冲应付"])}），'
                  f'差 {money(x["差额"])}　{x["说明"]}')
    print('== 项目（应收表 vs 导入） ==')
    for x in ch['项目应收核对']:
        print(f'  {x["结果"]} {x["项目"]}: 合同+补充 {x["合同加补充_导入"]}，开票 {x["开票_导入"]}，收款 {x["工程款_收支登记"]}')
    print('== 个人借款（公司欠） ==')
    for p, v in ch['个人借款（公司欠）'].items():
        print(f'  {v["结果"]} {p}: 收支登记 {money(v["收支登记"])}，个人明细 {money(v["个人明细"])}　{v["说明"]}')
    v = ch['考勤应发']
    print(f'== 考勤应发 == {v["结果"]} 复算 {money(v["按新规则复算"])}，1007 缓存 {money(v["原1007缓存"])}')
    print(f'== 日志 {len(ctx["log"])} 条、问题 {len(ctx["questions"])} 条，已写 _ctx.json ==')
    if '-v' in sys.argv:
        for x in ctx['log']:
            print('  ·', x)
        for x in ctx['questions']:
            print('  ?', x)


if __name__ == '__main__':
    main()
