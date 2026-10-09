# -*- coding: utf-8 -*-
"""查看表【工资表】【工资汇总】【用工成本表】【欠薪补发】（口径：报表口径.md §5）
   - 显示用「不带计提」的考勤数（考勤录了就显示）：考_应发（不含过账人员）、考_金额i（只给工程项目）、考_未分摊（办公室等公司管理）。
   - 已发＝收支登记归类「工资发放」、人员＝他（含农民工专户代发；工资退回是负数）。
   - 欠薪（某人到 d）＝期初欠薪（人员_期初欠薪＋期初「应付工资」）＋应发（月份≥考勤起算）－已发（日期≥考勤起算）。
     没有任何考勤记录的人（老板、按月发钱的管理人员）：发的钱直接算费用（跟 _收 的「冲工资」一样），不减欠薪。
     过账人员（姚俊强）单独标出，不算欠薪、不进合计。
   只用定义名称 ＋ 本模块几张表的格子；考勤的工种、补贴、扣款、备注按「表头锚定＋第 n 条」取（INDEX(考勤工资!$C$4:$C$20000,n+1)）。
   清单都用 counter/kth（隐藏列在可见区右边 AG 列以后；counter 的条件列从 CA 起）。"""
import datetime as dt
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation
from layout import *
from common import *

MONF = 'yyyy"年"m"月"'
DAYF = 'General;-General;""'                  # 天数：0 不显示
ACS = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
TEAM, EXTRA, DED, ANOTE = (inref(SH_ATT, c) for c in (ATT_TEAM, ATT_EXTRA, ATT_DED, ATT_NOTE))
CO_NAME = '公司管理（办公室等）'
# 全部人（不含过账）的期初欠薪
OP_ALL = 'SUMIFS(人员_期初欠薪,人员_过账,0)+SUMIFS(期初_金额,期初_类型,"应付工资")'
# 【工资汇总】隐藏列：收支登记逐条「冲欠薪的工资发放」（有考勤的人、不是过账、有人员）——全公司欠薪合计用（工资表也用）
SHOU_R0, SHOU_COL = 8, 'BP'
SHOU_FLAG = f"{SH_PAYSUM}!${SHOU_COL}${SHOU_R0}:${SHOU_COL}${SHOU_R0 + N_CASH - 1}"


# ───────────────────────── 小工具 ─────────────────────────
def hp(ws, cell, v):
    """隐藏辅助格"""
    ws[cell] = v
    ws[cell].font = F_HELP


def hide_range(ws, c0, c1):
    for i in range(CI(c0), CI(c1) + 1):
        ws.column_dimensions[CL(i)].hidden = True


def sum8(tpl):
    return '+'.join(tpl.format(i=i + 1) for i in range(NPAIR))


def renci(x, crit):
    """项目 x（转义后的格子）在 crit 条件下出现在几行考勤里（同一行出现几次只算一次）"""
    out = []
    for i in range(1, NPAIR + 1):
        c = [f'考_项目{i},{x}'] + list(crit) + [f'考_项目{j},"<>"&{x}' for j in range(1, i)]
        out.append('COUNTIFS(' + ','.join(c) + ')')
    return '+'.join(out)


def opening(e):
    return f'SUMIFS(人员_期初欠薪,人员_姓名,{e})+SUMIFS(期初_金额,期初_类型,"应付工资",期初_对象,{e})'


def owed(e, d, ym, has_att, op=None):
    """某人（e＝转义后的名字格）到 d 日（ym＝那个月 YYYYMM）的欠薪；has_att＝他有没有考勤记录（1/0 格）"""
    op = op or opening(e)
    return (f'{op}+SUMIFS(考_应发,考_姓名,{e},考_年月,">="&P_起算年月,考_年月,"<="&{ym})'
            f'+IF({has_att}=1,SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},收_日期,">="&P_考勤起算,收_日期,"<="&{d}),0)')


def owed_all(d, ym):
    """全公司（不含过账）到 d 日的欠薪合计"""
    return (f'{OP_ALL}+SUMIFS(考_应发,考_年月,">="&P_起算年月,考_年月,"<="&{ym})'
            f'+SUMIFS({SHOU_FLAG},收_日期,">="&P_考勤起算,收_日期,"<="&{d})')


def merge_put(ws, rng, v, font=F_TXT, fill_=None, fmt=None, align=AC, border=True):
    c0 = rng.split(':')[0]
    put(ws, c0, v, font, fill_, fmt, align, border)
    if ':' in rng:
        ws.merge_cells(rng)
        if border:
            a, b = rng.split(':')
            ca, ra = a.rstrip('0123456789'), int(a[len(a.rstrip('0123456789')):])
            cb, rb = b.rstrip('0123456789'), int(b[len(b.rstrip('0123456789')):])
            for r in range(ra, rb + 1):
                for c in range(CI(ca), CI(cb) + 1):
                    ws.cell(row=r, column=c).border = BD
                    if fill_ is not None:
                        ws.cell(row=r, column=c).fill = fill_


def sel_box(ws, lbl_cell, lbl, rng, fmt=None, dv=None, prompt=None, date=False, lbl_rng=None):
    """选择格：标签（可合并）＋亮黄输入格（可合并）"""
    c0 = rng.split(':')[0]
    selector(ws, lbl_cell, lbl, c0, None, dv_formula=dv, fmt=fmt, prompt=prompt)
    if lbl_rng:
        merge_put(ws, lbl_rng, None, F_KPI_L, fill('FFD9E1F2'))
    if ':' in rng:
        merge_put(ws, rng, None, F_SEL, FILL_SEL, fmt)
    if date:
        dv_date(ws, c0)


def note(ws, rng, v, align=AL, font=F_NOTE):
    merge_put(ws, rng, v, font, None, None, align, border=False)


def body(ws, r0, r1, c0, c1, key, fmts=None, aligns=None):
    """清单区：字体、对齐、格式；有内容的行才画框（条件格式），空着的行看不出来"""
    fmts, aligns = fmts or {}, aligns or {}
    for r in range(r0, r1 + 1):
        for i in range(CI(c0), CI(c1) + 1):
            c = CL(i)
            cell = ws[f'{c}{r}']
            cell.font = F_TXT
            cell.alignment = aligns.get(c, AC)
            if c in fmts:
                cell.number_format = fmts[c]
    ws.conditional_formatting.add(f'{c0}{r0}:{c1}{r1}', FormulaRule(formula=[f'${key}{r0}<>""'], border=BD))


def tot(ws, cells, fmts=None):
    """合计行：粉底加粗"""
    fmts = fmts or {}
    for coord, v in cells.items():
        col = coord.rstrip('0123456789')
        put(ws, coord, v, F_TXTB, FILL_TOT, fmts.get(col), AR if col in fmts else AC)


def consts(ws, col, r0, n):
    """隐藏常数列：第 r0 行起 1..n（各块「第几个」都从它取，公式能整列共享）"""
    for i in range(n):
        ws[f'{col}{r0 + i}'] = i + 1
        ws[f'{col}{r0 + i}'].font = F_HELP


# ═════════════════════════ 工资表（选一个月） ═════════════════════════
P_HDR, P_TOT, P_R0, P_N, P_NB = 5, 6, 7, 150, 20
P_PJC = [CL(4 + 2 * i) for i in range(NPAIR)]       # D F H J L N P R
P_DYC = [CL(5 + 2 * i) for i in range(NPAIR)]       # E G I K M O Q S
P_LAST = 'AF'


def build_pay(ws):
    R0, N = P_R0, P_N
    last = R0 + N - 1
    W = {'A': 5, 'B': 9, 'C': 8, 'T': 7, 'U': 9, 'V': 9, 'W': 8, 'X': 11, 'Y': 11, 'Z': 12, 'AA': 26, 'AB': 2,
         'AC': 16, 'AD': 6, 'AE': 7, 'AF': 12}
    for c in P_PJC:
        W[c] = 10
    for c in P_DYC:
        W[c] = 5
    widths(ws, W)
    tip = ('💡 黄格选月份（填这个月任意一天；空着＝首页截止日那个月）。列出这个月考勤的每一行：每个工地几天、应发多少；'
           '「本月已发」＝【收支登记】这个月发给他的工资（归类「工资发放」，含农民工专户代发，工资退回算负数）；「累计欠薪」＝到月底还欠他多少'
           '（建账前欠薪＋考勤起算月以来的应发－已发）。一个人超过 8 个工地占两行时，已发和欠薪只在第一行显示；这个月没考勤但发了工资的人排在最后。'
           '过账人员（如姚俊强）单独标出，不算欠薪、不进合计。右边是本月各项目人工。')
    title(ws, '工资表', P_LAST, C_VIEW, tip)
    home_link(ws, 'AG1')
    ws.column_dimensions['AG'].width = 10

    # ── 参数（隐藏 AR 列）──
    MS, ME, YM = '$AR$3', '$AR$4', '$AR$5'
    hp(ws, 'AR3', '=IF(ISNUMBER(C3),DATE(YEAR(C3),MONTH(C3),1),DATE(YEAR(P_截止),MONTH(P_截止),1))')
    hp(ws, 'AR4', '=DATE(YEAR(AR3),MONTH(AR3)+1,0)')
    hp(ws, 'AR5', '=YEAR(AR3)*100+MONTH(AR3)')
    CRIT = [f'考_年月,{YM}', '考_过账,0']
    ws['AR3'].number_format = ws['AR4'].number_format = DATE

    # ── 选择格、实际值、核对 ──
    sel_box(ws, 'B3', '选月份', 'C3:D3', fmt=MONF, date=True)
    note(ws, 'E3:K3', f'="实际："&YEAR({MS})&"年"&MONTH({MS})&"月"&IF(ISNUMBER(C3),"","（空着＝首页截止日那个月）")',
         font=F_AUTOB)
    merge_put(ws, 'T3:W3', '收支登记本月发工资（不含过账）', F_KPI_L, fill('FFD9E1F2'))
    put(ws, 'X3', f'=-SUMIFS(收_净额,收_归类,"工资发放",收_过账,0,收_年月,{YM})', F_AUTOB, FILL_AUTO, MONEY, AR)
    merge_put(ws, 'T4:W4', '其中下面没列到的（没填人员等）', F_KPI_L, fill('FFD9E1F2'))
    put(ws, 'X4', f'=ROUND(X3-Y{P_TOT},2)', F_AUTOB, FILL_AUTO, MONEY, AR)
    merge_put(ws, 'Y3:Y4', '全公司到月底欠薪', F_KPI_L, fill('FFD9E1F2'), align=ACW)
    merge_put(ws, 'Z3:Z4', f'=ROUND({owed_all(ME, YM)},2)', F_AUTOB, FILL_AUTO, MONEY, AR)

    # ── 隐藏：常数列 AH（第 R0 行起 1..N_ATT）──
    consts(ws, 'AH', R0, N_ATT)
    K = lambda r: f'$AH{r}'
    # 本月考勤行（_考 第 i 条）
    counter(ws, 'AT', R0, N_ATT, lambda i: f'AND(INDEX(考_年月,$AH{R0 + i})={YM},INDEX(考_有效,$AH{R0 + i})=1)')
    K1 = cnt('AT', R0, N_ATT)
    # 本月没考勤、但有工资发放的人（人员信息第 i 个）
    for i in range(N_PER):
        r = R0 + i
        hp(ws, f'AV{r}', f'=INDEX(人员_姓名,$AH{r})&""')
        hp(ws, f'AW{r}', f'=IF(AV{r}="","",{esc(f"AV{r}")})')
    counter(ws, 'AY', R0, N_PER, lambda i: (
        f'AND(AV{R0 + i}<>"",COUNTIFS(考_姓名,AW{R0 + i},考_年月,{YM},考_有效,1)=0,'
        f'SUMIFS(收_净额,收_归类,"工资发放",收_人员,AW{R0 + i},收_年月,{YM})<>0)'))
    K2 = cnt('AY', R0, N_PER)
    # 本月有人工的工程项目（项目档案第 i 个）
    for i in range(N_PJ):
        r = R0 + i
        hp(ws, f'BA{r}', f'=INDEX(项目_名称,$AH{r})&""')
        hp(ws, f'BB{r}', f'=IF(BA{r}="","",{esc(f"BA{r}")})')
        skip = f'OR(BA{r}="",INDEX(项目_类型,$AH{r})<>"工程")'
        hp(ws, f'BC{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_天数{{i}},考_项目{{i}},BB{r},考_年月,{YM},考_过账,0)') + ')')
        hp(ws, f'BD{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_金额{{i}},考_项目{{i}},BB{r},考_年月,{YM})') + ')')
    counter(ws, 'BE', R0, N_PJ, lambda i: f'OR(BC{R0 + i}<>0,BD{R0 + i}<>0)')
    KP = cnt('BE', R0, N_PJ)

    # ── 表头、合计 ──
    heads = [('A', '序号'), ('B', '姓名'), ('C', '工种/类别')]
    for i in range(NPAIR):
        heads += [(P_PJC[i], f'项目{i + 1}'), (P_DYC[i], f'天数{i + 1}')]
    heads += [('T', '总天数'), ('U', '单价\n（月薪的是月薪）'), ('V', '加班补贴'), ('W', '扣款'), ('X', '应发'),
              ('Y', '本月已发'), ('Z', '累计欠薪\n（到月底）'), ('AA', '备注')]
    header(ws, P_HDR, heads, C_VIEW, height=36)
    note(ws, f'A4:R4', f'="共 "&({K1}+{K2})&" 行：本月考勤 "&{K1}&" 行"&IF({K2}>0,"，本月没考勤、只有发放 "&{K2}&" 人","")'
                       f'&IF({K1}+{K2}>{N},"；只显示前 {N} 行","")')
    rt = P_TOT
    rng = lambda c: f'{c}{R0}:{c}{last}'
    merge_put(ws, f'A{rt}:C{rt}', f'="合计（"&COUNTIFS({rng("$AM")},1,{rng("$AN")},0)&" 人）"', F_TXTB, FILL_TOT)
    tc = {f'{c}{rt}': None for c in P_PJC + P_DYC + ['U']}
    for c in ('T', 'V', 'W', 'X', 'Y'):
        tc[f'{c}{rt}'] = f'=SUMIFS({rng(c)},{rng("$AN")},0)'
    tc[f'Z{rt}'] = f'=SUM({rng("Z")})'
    tot(ws, tc, {'T': DAYF, 'V': MONEY, 'W': MONEY, 'X': MONEY, 'Y': MONEY, 'Z': MONEY})
    put(ws, f'AA{rt}', '合计不含过账人员', F_NOTE, FILL_TOT, align=AL)

    # ── 清单 ──
    for i in range(N):
        r = R0 + i
        k = K(r)
        hp(ws, f'AI{r}', '=' + kth(k, 'AT', R0, N_ATT))
        hp(ws, f'AJ{r}', f'=IF(AI{r}>0,0,IF({k}-{K1}<1,0,{kth(f"({k}-{K1})", "AY", R0, N_PER)}))')
        hp(ws, f'AK{r}', f'=IF(AI{r}>0,INDEX(考_姓名,AI{r}),IF(AJ{r}>0,INDEX(人员_姓名,AJ{r}),""))&""')
        hp(ws, f'AL{r}', f'=IF(AK{r}="","",{esc(f"AK{r}")})')
        hp(ws, f'AM{r}', f'=IF(AI{r}>0,IF(COUNTIFS(考_姓名,AL{r},考_年月,{YM},考_有效,1,考_n,"<"&AI{r})=0,1,0),IF(AJ{r}>0,1,0))')
        hp(ws, f'AN{r}', f'=IF(AI{r}>0,INDEX(考_过账,AI{r}),IF(AJ{r}>0,INDEX(人员_过账,AJ{r}),0))')
        hp(ws, f'AO{r}', f'=IF(AK{r}="",0,IF(COUNTIF(考_姓名,AL{r})>0,1,0))')
        hp(ws, f'AP{r}', f'=IF(AK{r}="","",IF(AN{r}=1,"过账人员：只在公司账上过一下，不算工资、不算欠薪",'
                         f'IF(AJ{r}>0,IF(AO{r}=0,"没有考勤记录：发的工资直接算费用，不算欠薪","本月没考勤，只有发放"),'
                         f'IF(AM{r}=0,"同一人第 2 行：已发、欠薪看他上一行",""))))')
        ws[f'A{r}'] = f'=IF(AK{r}="","",{k})'
        ws[f'B{r}'] = f'=AK{r}'
        ws[f'C{r}'] = (f'=IF(AI{r}>0,IF(INDEX({TEAM},AI{r}+1)&""<>"",INDEX({TEAM},AI{r}+1)&"",'
                       f'IFERROR(INDEX(人员_类别,MATCH(AK{r},人员_姓名,0))&"","")),IF(AJ{r}>0,INDEX(人员_类别,AJ{r})&"",""))')
        for j in range(NPAIR):
            ws[f'{P_PJC[j]}{r}'] = f'=IF($AI{r}=0,"",INDEX(考_项目{j + 1},$AI{r})&"")'
            ws[f'{P_DYC[j]}{r}'] = f'=IF($AI{r}=0,"",IF(INDEX(考_天数{j + 1},$AI{r})=0,"",INDEX(考_天数{j + 1},$AI{r})))'
        ws[f'T{r}'] = f'=IF(AI{r}=0,"",INDEX(考_总天数,AI{r}))'
        ws[f'U{r}'] = f'=IF(AI{r}=0,"",INDEX(考_单价,AI{r}))'
        ws[f'V{r}'] = f'=IF(AI{r}=0,"",N(INDEX({EXTRA},AI{r}+1)))'
        ws[f'W{r}'] = f'=IF(AI{r}=0,"",N(INDEX({DED},AI{r}+1)))'
        ws[f'X{r}'] = f'=IF(AI{r}>0,IF(AN{r}=1,INDEX(考_过账应发,AI{r}),INDEX(考_应发,AI{r})),IF(AJ{r}>0,0,""))'
        ws[f'Y{r}'] = f'=IF(AM{r}=1,-SUMIFS(收_净额,收_归类,"工资发放",收_人员,AL{r},收_年月,{YM}),"")'
        ws[f'Z{r}'] = f'=IF(AND(AM{r}=1,AN{r}=0),ROUND({owed(f"AL{r}", ME, YM, f"AO{r}")},2),"")'
        an = f'IF(AI{r}>0,INDEX({ANOTE},AI{r}+1)&"","")'
        ws[f'AA{r}'] = f'=IF(AP{r}="",{an},IF({an}="",AP{r},AP{r}&"；"&{an}))'
    fm = {c: DAYF for c in P_DYC + ['T']}
    fm.update({'U': MONEY, 'V': MONEY, 'W': MONEY, 'X': MONEY, 'Y': MONEY, 'Z': MONEY})
    al = {c: ACS for c in P_PJC + ['B', 'C']}
    al['AA'] = AL
    body(ws, R0, last, 'A', 'AA', 'B', fm, al)
    # 过账人员那行灰字
    ws.conditional_formatting.add(f'A{R0}:AA{last}', FormulaRule(formula=[f'$AN{R0}=1'], font=Font(name=YH, sz=10, color='FF808080', italic=True)))

    # ── 右边：本月各项目人工 ──
    section(ws, 4, 'AC', 'AF', '本月各项目人工', C_VIEW)
    header(ws, P_HDR, [('AC', '项目'), ('AD', '人数'), ('AE', '天数'), ('AF', '人工金额\n（应发按天数分）')], C_VIEW, height=36)
    tot(ws, {f'AC{rt}': '合计', f'AD{rt}': f'=SUM(AD{R0}:AD{R0 + P_NB - 1})', f'AE{rt}': f'=SUM(AE{R0}:AE{R0 + P_NB - 1})',
             f'AF{rt}': f'=SUM(AF{R0}:AF{R0 + P_NB - 1})'}, {'AD': INT, 'AE': DAYF, 'AF': MONEY})
    for i in range(P_NB):
        r = R0 + i
        k = K(r)
        hp(ws, f'BG{r}', '=' + kth(k, 'BE', R0, N_PJ))
        hp(ws, f'BH{r}', f'=IF(BG{r}=0,"",INDEX($BB${R0}:$BB${R0 + N_PJ - 1},BG{r}))')
        co = f'{k}={KP}+1'
        ws[f'AC{r}'] = f'=IF(BG{r}>0,INDEX($BA${R0}:$BA${R0 + N_PJ - 1},BG{r}),IF({co},"{CO_NAME}",""))'
        ws[f'AD{r}'] = (f'=IF(BG{r}>0,{renci(f"BH{r}", CRIT)},'
                        f'IF({co},COUNTIFS(考_年月,{YM},考_未分摊,"<>0"),""))')
        ws[f'AE{r}'] = (f'=IF(BG{r}>0,INDEX($BC${R0}:$BC${R0 + N_PJ - 1},BG{r}),'
                        f'IF({co},SUMIFS(考_总天数,考_年月,{YM},考_过账,0)-SUM($BC${R0}:$BC${R0 + N_PJ - 1}),""))')
        ws[f'AF{r}'] = (f'=IF(BG{r}>0,INDEX($BD${R0}:$BD${R0 + N_PJ - 1},BG{r}),'
                        f'IF({co},SUMIFS(考_未分摊,考_年月,{YM}),""))')
    body(ws, R0, R0 + P_NB - 1, 'AC', 'AF', 'AC', {'AD': INT, 'AE': DAYF, 'AF': MONEY}, {'AC': AL})
    note(ws, f'AC{R0 + P_NB}:AF{R0 + P_NB}', f'=IF({KP}+1>{P_NB},"共 "&{KP}&" 个项目，只显示前 {P_NB - 1} 个","")')

    hide_range(ws, 'AH', 'BZ')
    for c in ('AT', 'AY', 'BE'):
        ws.column_dimensions[c].hidden = True
    ws.freeze_panes = f'D{R0}'
    print_setup(ws, f'{P_HDR}:{P_TOT}', landscape=True)
    ws.print_area = f'A1:{P_LAST}{last}'


# ═════════════════════════ 工资汇总（起止月份） ═════════════════════════
S_SEC, S_HDR, S_TOT, S_R0, S_N = 5, 6, 7, 8, 300
S_NB = 20                                    # ② 项目最多显示几行（含「公司管理」）
S_MSEC, S_MHDR, S_MTOT, S_MR0, S_MN = 29, 30, 31, 32, 60
S_LAST = 'N'


def build_paysum(ws):
    R0, N = S_R0, S_N
    last = R0 + N - 1
    widths(ws, {'A': 5, 'B': 9, 'C': 8, 'D': 12, 'E': 8, 'F': 12, 'G': 12, 'H': 12, 'I': 30, 'J': 2,
                'K': 18, 'L': 12, 'M': 12, 'N': 12})
    tip = ('💡 黄格选起止月份（填那个月任意一天；空着＝首页年初～截止日那个月）。① 每人：期初欠薪（起始月以前累计欠他的）＋应发－已发＝期末欠薪，'
           '只列这段时间有考勤、有发放或者有欠薪的人；② 按项目：人工金额＝考勤应发按天数分到各工地，在办公室的算「公司管理」，合计＝应发合计；'
           '③ 每月应发、已发、月底欠薪合计。欠薪从考勤起算月开始算（以前发的工资直接算成本）；没有考勤记录的人（老板、按月发钱的管理人员）发的钱直接算费用，'
           '不算欠薪；过账人员不算工资。')
    title(ws, '工资汇总', S_LAST, C_VIEW, tip)
    home_link(ws, 'O1')
    ws.column_dimensions['O'].width = 10

    # ── 参数（隐藏 AJ 列）──
    MS0, ME1, YM0, YM1 = '$AJ$3', '$AJ$4', '$AJ$5', '$AJ$6'
    OPA, GF, GG, GH = '$AJ$7', '$AJ$8', '$AJ$9', '$AJ$10'
    hp(ws, 'AJ3', '=IF(ISNUMBER(C3),DATE(YEAR(C3),MONTH(C3),1),P_年初)')
    hp(ws, 'AJ4', '=IF(ISNUMBER(C4),DATE(YEAR(C4),MONTH(C4)+1,0),DATE(YEAR(P_截止),MONTH(P_截止)+1,0))')
    hp(ws, 'AJ5', '=YEAR(AJ3)*100+MONTH(AJ3)')
    hp(ws, 'AJ6', '=YEAR(AJ4)*100+MONTH(AJ4)')
    hp(ws, 'AJ7', f'={OP_ALL}')
    RNG_YM = f'考_年月,">="&{YM0},考_年月,"<="&{YM1}'
    CRIT = [f'考_年月,">="&{YM0}', f'考_年月,"<="&{YM1}', '考_过账,0']
    RNG_D = f'收_日期,">="&{MS0},收_日期,"<="&{ME1}'
    hp(ws, 'AJ8', f'=SUMIFS(考_应发,{RNG_YM})')
    hp(ws, 'AJ9', f'=-SUMIFS(收_净额,收_归类,"工资发放",收_过账,0,{RNG_D})')
    hp(ws, 'AJ10', f'=ROUND({owed_all(ME1, YM1)},2)')
    for c in ('AJ3', 'AJ4'):
        ws[c].number_format = DATE

    # ── 选择格 ──
    sel_box(ws, 'B3', '起始月份', 'C3:D3', fmt=MONF, date=True)
    sel_box(ws, 'B4', '结束月份', 'C4:D4', fmt=MONF, date=True)
    note(ws, 'E3:I3', f'="实际："&YEAR({MS0})&"年"&MONTH({MS0})&"月 ～ "&YEAR({ME1})&"年"&MONTH({ME1})&"月"'
                      f'&IF(AND(ISNUMBER(C3),ISNUMBER(C4)),"","（空着＝首页年初～截止日那个月）")&IF({ME1}<{MS0},"　⚠ 结束早于起始","")',
         font=F_AUTOB)
    note(ws, 'E4:I4', '="欠薪从考勤起算月（"&YEAR(P_考勤起算)&"年"&MONTH(P_考勤起算)&"月）开始算：建账前欠薪＋应发－已发"')

    consts(ws, 'AH', R0, N_CASH)
    K = lambda r: f'$AH{r}'

    # ── 隐藏：每个人（人员信息第 i 个，第 R0 行起）──
    PR = lambda c: f'${c}${R0}:${c}${R0 + N_PER - 1}'
    for i in range(N_PER):
        r = R0 + i
        k = K(r)
        nm, e, gz, att, op = f'AL{r}', f'AM{r}', f'AN{r}', f'AO{r}', f'AP{r}'
        hp(ws, nm, f'=INDEX(人员_姓名,{k})&""')
        hp(ws, e, f'=IF({nm}="","",{esc(nm)})')
        hp(ws, gz, f'=IF({nm}="",0,INDEX(人员_过账,{k}))')
        hp(ws, att, f'=IF({nm}="",0,IF(COUNTIF(考_姓名,{e})>0,1,0))')
        hp(ws, op, f'=IF({nm}="",0,INDEX(人员_期初欠薪,{k})+SUMIFS(期初_金额,期初_类型,"应付工资",期初_对象,{e}))')
        hp(ws, f'AQ{r}', f'=IF(OR({nm}="",{gz}=1),0,ROUND({op}+SUMIFS(考_应发,考_姓名,{e},考_年月,">="&P_起算年月,考_年月,"<"&{YM0})'
                         f'+IF({att}=1,SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},收_日期,">="&P_考勤起算,收_日期,"<"&{MS0}),0),2))')
        hp(ws, f'AR{r}', f'=IF({nm}="",0,SUMIFS(考_总天数,考_姓名,{e},{RNG_YM},考_有效,1))')
        hp(ws, f'AS{r}', f'=IF({nm}="",0,IF({gz}=1,SUMIFS(考_过账应发,考_姓名,{e},{RNG_YM}),SUMIFS(考_应发,考_姓名,{e},{RNG_YM})))')
        hp(ws, f'AT{r}', f'=IF({nm}="",0,-SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},{RNG_D}))')
        hp(ws, f'AU{r}', f'=IF(OR({nm}="",{gz}=1),0,ROUND({owed(e, ME1, YM1, att, op)},2))')
        hp(ws, f'AV{r}', f'=IF(OR({nm}="",{MS0}>=P_考勤起算),0,IF(OR(SUMIFS(考_应发,考_姓名,{e},考_年月,">="&{YM0},考_年月,"<"&P_起算年月)<>0,'
                         f'SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},收_日期,">="&{MS0},收_日期,"<"&P_考勤起算)<>0),1,0))')
    counter(ws, 'AW', R0, N_PER, lambda i: (f'AND(AL{R0 + i}<>"",OR(AR{R0 + i}<>0,AS{R0 + i}<>0,AT{R0 + i}<>0,'
                                             f'AQ{R0 + i}<>0,AU{R0 + i}<>0))'))
    KN = cnt('AW', R0, N_PER)

    # ── ① 按人 ──
    section(ws, S_SEC, 'A', 'I', '① 按人', C_VIEW)
    ws[f'A{S_SEC}'] = f'="① 按人（共 "&{KN}&" 人"&IF({KN}>{N},"，只显示前 {N} 人","")&"）"'
    header(ws, S_HDR, [('A', '序号'), ('B', '姓名'), ('C', '类别'), ('D', '期初欠薪\n（起始月以前）'), ('E', '天数'), ('F', '应发'),
                       ('G', '已发'), ('H', '期末欠薪'), ('I', '备注')], C_VIEW, height=36)
    rt = S_TOT
    rng = lambda c: f'{c}{R0}:{c}{last}'
    merge_put(ws, f'A{rt}:C{rt}', f'="合计（"&COUNTIFS({rng("$AZ")},0,{rng("$AY")},">0")&" 人）"', F_TXTB, FILL_TOT)
    tot(ws, {f'D{rt}': f'=SUM({rng("D")})', f'E{rt}': f'=SUMIFS({rng("E")},{rng("$AZ")},0)',
             f'F{rt}': f'=SUMIFS({rng("F")},{rng("$AZ")},0)', f'G{rt}': f'=SUMIFS({rng("G")},{rng("$AZ")},0)',
             f'H{rt}': f'=SUM({rng("H")})'}, {'D': MONEY, 'E': DAYF, 'F': MONEY, 'G': MONEY, 'H': MONEY})
    put(ws, f'I{rt}', f'=IF(AND(ROUND(F{rt}-{GF},0)=0,ROUND(G{rt}-{GG},0)=0,ROUND(H{rt}-{GH},0)=0),"不含过账人员",'
                      f'"⚠ 全部考勤/发放：应发 "&TEXT({GF},"#,##0")&"、已发 "&TEXT({GG},"#,##0")&"、期末欠薪 "&TEXT({GH},"#,##0")'
                      f'&"（差额＝名字不在【人员信息】里，或工资发放没填人员）")', F_NOTE, FILL_TOT, align=ALW)
    for i in range(N):
        r = R0 + i
        k = K(r)
        p = f'AY{r}'
        hp(ws, p, '=' + kth(k, 'AW', R0, N_PER))
        hp(ws, f'AZ{r}', f'=IF({p}=0,0,INDEX({PR("AN")},{p}))')
        g = lambda c: f'INDEX({PR(c)},{p})'
        ws[f'A{r}'] = f'=IF({p}=0,"",{k})'
        ws[f'B{r}'] = f'=IF({p}=0,"",{g("AL")})'
        ws[f'C{r}'] = f'=IF({p}=0,"",INDEX(人员_类别,{p})&"")'
        ws[f'D{r}'] = f'=IF({p}=0,"",IF(AZ{r}=1,"",{g("AQ")}))'
        ws[f'E{r}'] = f'=IF({p}=0,"",{g("AR")})'
        ws[f'F{r}'] = f'=IF({p}=0,"",{g("AS")})'
        ws[f'G{r}'] = f'=IF({p}=0,"",{g("AT")})'
        ws[f'H{r}'] = f'=IF({p}=0,"",IF(AZ{r}=1,"",{g("AU")}))'
        ws[f'I{r}'] = (f'=IF({p}=0,"",IF(AZ{r}=1,"过账人员：不算工资、不算欠薪（不进合计）",IF({g("AO")}=0,"没有考勤记录：发的工资直接算费用，不算欠薪",'
                       f'IF({g("AV")}=1,"考勤起算月以前的应发、已发不算欠薪",""))))')
    body(ws, R0, last, 'A', 'I', 'B', {'D': MONEY, 'E': DAYF, 'F': MONEY, 'G': MONEY, 'H': MONEY}, {'I': AL})
    ws.conditional_formatting.add(f'A{R0}:I{last}', FormulaRule(formula=[f'$AZ{R0}=1'], font=Font(name=YH, sz=10, color='FF808080', italic=True)))

    # ── ② 按项目（右边）──
    for i in range(N_PJ):
        r = R0 + i
        k = K(r)
        hp(ws, f'BB{r}', f'=INDEX(项目_名称,{k})&""')
        hp(ws, f'BC{r}', f'=IF(BB{r}="","",{esc(f"BB{r}")})')
        skip = f'OR(BB{r}="",INDEX(项目_类型,{k})<>"工程")'
        hp(ws, f'BD{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_天数{{i}},考_项目{{i}},BC{r},{RNG_YM},考_过账,0)') + ')')
        hp(ws, f'BE{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_金额{{i}},考_项目{{i}},BC{r},{RNG_YM})') + ')')
    counter(ws, 'BF', R0, N_PJ, lambda i: f'OR(BD{R0 + i}<>0,BE{R0 + i}<>0)')
    KP = cnt('BF', R0, N_PJ)
    PJ = lambda c: f'${c}${R0}:${c}${R0 + N_PJ - 1}'
    section(ws, S_SEC, 'K', 'N', '② 按项目', C_VIEW)
    ws[f'K{S_SEC}'] = f'="② 按项目"&IF({KP}+1>{S_NB},"（共 "&{KP}&" 个，只显示前 {S_NB - 1} 个）","")'
    header(ws, S_HDR, [('K', '项目'), ('L', '人次'), ('M', '天数'), ('N', '人工金额')], C_VIEW, height=36)
    b1 = R0 + S_NB - 1
    tot(ws, {f'K{rt}': '合计（＝应发合计）', f'L{rt}': f'=SUM(L{R0}:L{b1})', f'M{rt}': f'=SUM(M{R0}:M{b1})', f'N{rt}': f'=SUM(N{R0}:N{b1})'},
        {'L': INT, 'M': DAYF, 'N': MONEY})
    for i in range(S_NB):
        r = R0 + i
        k = K(r)
        hp(ws, f'BH{r}', '=' + kth(k, 'BF', R0, N_PJ))
        hp(ws, f'BI{r}', f'=IF(BH{r}=0,"",INDEX({PJ("BC")},BH{r}))')
        co = f'{k}={KP}+1'
        ws[f'K{r}'] = f'=IF(BH{r}>0,INDEX({PJ("BB")},BH{r}),IF({co},"{CO_NAME}",""))'
        ws[f'L{r}'] = (f'=IF(BH{r}>0,{renci(f"BI{r}", CRIT)},'
                       f'IF({co},COUNTIFS({RNG_YM},考_未分摊,"<>0"),""))')
        ws[f'M{r}'] = (f'=IF(BH{r}>0,INDEX({PJ("BD")},BH{r}),IF({co},SUMIFS(考_总天数,{RNG_YM},考_过账,0)-SUM({PJ("BD")}),""))')
        ws[f'N{r}'] = f'=IF(BH{r}>0,INDEX({PJ("BE")},BH{r}),IF({co},SUMIFS(考_未分摊,{RNG_YM}),""))'
    body(ws, R0, b1, 'K', 'N', 'K', {'L': INT, 'M': DAYF, 'N': MONEY}, {'K': AL})

    # ── ③ 按月（右边，② 下面）──
    section(ws, S_MSEC, 'K', 'N', '③ 按月', C_VIEW)
    NMON = f'((YEAR({ME1})-YEAR({MS0}))*12+MONTH({ME1})-MONTH({MS0})+1)'
    ws[f'K{S_MSEC}'] = f'="③ 按月（"&MAX(0,{NMON})&" 个月"&IF({NMON}>{S_MN},"，只显示前 {S_MN} 个","")&"；已发不含过账）"'
    header(ws, S_MHDR, [('K', '月份'), ('L', '应发'), ('M', '已发'), ('N', '月底欠薪')], C_VIEW, height=None)
    m1 = S_MR0 + S_MN - 1
    tot(ws, {f'K{S_MTOT}': '合计（欠薪＝期末）', f'L{S_MTOT}': f'=SUM(L{S_MR0}:L{m1})', f'M{S_MTOT}': f'=SUM(M{S_MR0}:M{m1})',
             f'N{S_MTOT}': f'={GH}'}, {'L': MONEY, 'M': MONEY, 'N': MONEY})
    for i in range(S_MN):
        r = S_MR0 + i
        hp(ws, f'BK{r}', f'=DATE(YEAR({MS0}),MONTH({MS0})+{i},1)')
        hp(ws, f'BL{r}', f'=DATE(YEAR(BK{r}),MONTH(BK{r})+1,0)')
        hp(ws, f'BM{r}', f'=YEAR(BK{r})*100+MONTH(BK{r})')
        ok = f'BK{r}>{ME1}'
        ws[f'K{r}'] = f'=IF({ok},"",BK{r})'
        ws[f'L{r}'] = f'=IF({ok},"",SUMIFS(考_应发,考_年月,BM{r}))'
        ws[f'M{r}'] = f'=IF({ok},"",-SUMIFS(收_净额,收_归类,"工资发放",收_过账,0,收_年月,BM{r}))'
        ws[f'N{r}'] = f'=IF({ok},"",ROUND({owed_all(f"BL{r}", f"BM{r}")},2))'
    body(ws, S_MR0, m1, 'K', 'N', 'K', {'K': MONF, 'L': MONEY, 'M': MONEY, 'N': MONEY})

    # ── 隐藏：收支登记逐条「冲欠薪的工资发放」（有考勤的人；全公司欠薪合计用）──
    for i in range(N_CASH):
        r = SHOU_R0 + i
        n = K(r)
        who = f'INDEX(收_人员,{n})'
        hp(ws, f'{SHOU_COL}{r}', f'=IF(INDEX(收_归类,{n})<>"工资发放",0,IF(INDEX(收_过账,{n})=1,0,IF({who}="",0,'
                                 f'IF(COUNTIF(考_姓名,{esc(who)})>0,INDEX(收_净额,{n}),0))))')

    hide_range(ws, 'AH', 'BZ')
    for c in ('AW', 'BF'):
        ws.column_dimensions[c].hidden = True
    ws.freeze_panes = f'C{R0}'
    print_setup(ws, f'{S_HDR}:{S_TOT}', landscape=True)
    ws.print_area = f'A1:{S_LAST}{last}'


# ═════════════════════════ 用工成本表（起止月份＋最多 8 个工地） ═════════════════════════
L_SEC, L_H1, L_H2, L_TOT, L_R0 = 7, 8, 9, 10, 11
L_N, L_NS = 200, 40                          # ② 最多几个人；① 最多几个工地（含「公司管理」）
L_DC = [CL(3 + 2 * i) for i in range(NPAIR)]   # ② 天数列 C E G I K M O Q
L_AC = [CL(4 + 2 * i) for i in range(NPAIR)]   # ② 金额列 D F H J L N P R
L_LAST = 'Z'


def build_labor(ws):
    R0 = L_R0
    last = R0 + L_N - 1
    W = {'A': 5, 'B': 10, 'S': 8, 'T': 12, 'U': 2, 'V': 18, 'W': 9, 'X': 12, 'Y': 7, 'Z': 10}
    for c in L_DC:
        W[c] = 7
    for c in L_AC:
        W[c] = 11
    widths(ws, W)
    tip = ('💡 黄格选起止月份（填那个月任意一天；空着＝首页年初～截止日那个月），再选工地（最多 8 个，下拉选；8 个都空＝全部工程项目）。'
           '右边 ① 每个工地的天数、人工成本（考勤应发按天数分到工地）、人次、平均每天；左边 ② 每个人在每个工地干了几天、多少钱。'
           '没选工地时 ① 按人工成本从大到小列全部工地，另加一行「公司管理」（办公室等，不算工地），② 只列人工成本最大的 8 个工地。过账人员不算。')
    title(ws, '用工成本表', L_LAST, C_VIEW, tip)
    home_link(ws, 'AA1')
    ws.column_dimensions['AA'].width = 10

    # ── 参数（隐藏 AH 列）──
    MS0, ME1, YM0, YM1, KPA = '$AH$3', '$AH$4', '$AH$5', '$AH$6', '$AH$7'
    hp(ws, 'AH3', '=IF(ISNUMBER(C3),DATE(YEAR(C3),MONTH(C3),1),P_年初)')
    hp(ws, 'AH4', '=IF(ISNUMBER(G3),DATE(YEAR(G3),MONTH(G3)+1,0),DATE(YEAR(P_截止),MONTH(P_截止)+1,0))')
    hp(ws, 'AH5', '=YEAR(AH3)*100+MONTH(AH3)')
    hp(ws, 'AH6', '=YEAR(AH4)*100+MONTH(AH4)')
    for c in ('AH3', 'AH4'):
        ws[c].number_format = DATE
    RNG = f'考_年月,">="&{YM0},考_年月,"<="&{YM1}'
    CRIT = [f'考_年月,">="&{YM0}', f'考_年月,"<="&{YM1}', '考_过账,0']

    # ── 选择格 ──
    sel_box(ws, 'B3', '起始月份', 'C3:D3', fmt=MONF, date=True)
    sel_box(ws, 'E3', '结束月份', 'G3:H3', fmt=MONF, date=True, lbl_rng='E3:F3')
    note(ws, 'I3:T3', f'="实际："&YEAR({MS0})&"年"&MONTH({MS0})&"月 ～ "&YEAR({ME1})&"年"&MONTH({ME1})&"月"'
                      f'&IF(AND(ISNUMBER(C3),ISNUMBER(G3)),"","（空着＝首页年初～截止日那个月）")&IF({ME1}<{MS0},"　⚠ 结束早于起始","")',
         font=F_AUTOB)
    put(ws, 'B4', '选工地', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B5', '（下拉）', F_NOTE, align=AC)
    sels = []
    for k in range(NPAIR):
        d, a = L_DC[k], L_AC[k]
        merge_put(ws, f'{d}4:{a}4', f'工地{k + 1}', F_KPI_L, fill('FFD9E1F2'))
        merge_put(ws, f'{d}5:{a}5', None, F_SEL, FILL_SEL)
        sels.append(f'{d}5')
    dv = DataValidation(type='list', formula1='=项目列表', allow_blank=True, showErrorMessage=False)
    dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', '选工地；8 个都空＝全部工程项目', True
    ws.add_data_validation(dv)
    for c in sels:
        dv.add(c)
    note(ws, 'S4:T5', '8 个都空＝全部工地', align=ACW)

    # 所选工地去重、挤掉空格（AJ 名字、AK 是不是新的、AL 累计）
    for k in range(NPAIR):
        r = 3 + k
        hp(ws, f'AJ{r}', f'=TRIM({sels[k]}&"")')
        hp(ws, f'AK{r}', '=IF(AJ3="",0,1)' if k == 0 else f'=IF(AJ{r}="",0,IF(COUNTIF($AJ$3:AJ{r - 1},{esc(f"AJ{r}")})=0,1,0))')
        hp(ws, f'AL{r}', f'=SUM($AK$3:AK{r})')
    KSEL = '$AL$10'

    consts(ws, 'AG', R0, N_ATT)
    K = lambda r: f'$AG{r}'
    # ── 隐藏：项目（项目档案第 i 个）——没选工地时用：天数、人工、排序键（人工×100×1000＋(1000－i)，LARGE 取大的）──
    for i in range(N_PJ):
        r = R0 + i
        k = K(r)
        hp(ws, f'AO{r}', f'=INDEX(项目_名称,{k})&""')
        hp(ws, f'AP{r}', f'=IF(AO{r}="","",{esc(f"AO{r}")})')
        hp(ws, f'AQ{r}', f'=IF(AO{r}="","",INDEX(项目_类型,{k})&"")')
        skip = f'OR(AQ{r}<>"工程",{KSEL}>0)'
        hp(ws, f'AR{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_天数{{i}},考_项目{{i}},AP{r},{RNG},考_过账,0)') + ')')
        hp(ws, f'AS{r}', f'=IF({skip},0,' + sum8(f'SUMIFS(考_金额{{i}},考_项目{{i}},AP{r},{RNG})') + ')')
        hp(ws, f'AU{r}', f'=IF(AND(AQ{r}="工程",OR(AR{r}<>0,AS{r}<>0)),ROUND(AS{r}*100,0)*1000+1000-{k},"")')
    PJ = lambda c: f'${c}${R0}:${c}${R0 + N_PJ - 1}'
    hp(ws, 'AH7', f'=COUNT({PJ("AU")})')

    # ── 右边 ① 按工地 ──
    s1 = R0 + L_NS - 1
    section(ws, L_SEC, 'V', 'Z', '① 按工地', C_VIEW)
    ws[f'V{L_SEC}'] = (f'="① 按工地"&IF({KSEL}>0,"（选了 "&{KSEL}&" 个）","（没选＝全部 "&{KPA}&" 个工地）")'
                       f'&IF(AND({KSEL}=0,{KPA}+1>{L_NS}),"只显示前 {L_NS - 1} 个","")')
    for c, t in (('V', '工地'), ('W', '天数'), ('X', '人工成本'), ('Y', '人次'), ('Z', '平均每天')):
        merge_put(ws, f'{c}{L_H1}:{c}{L_H2}', t, F_HDR, fill(C_VIEW), align=ACW)
    rt = L_TOT
    tot(ws, {f'V{rt}': '合计', f'W{rt}': f'=SUM(W{R0}:W{s1})', f'X{rt}': f'=SUM(X{R0}:X{s1})', f'Y{rt}': f'=SUM(Y{R0}:Y{s1})',
             f'Z{rt}': f'=IF(N(W{rt})=0,"",X{rt}/W{rt})'}, {'W': DAYF, 'X': MONEY, 'Y': INT, 'Z': MONEY})
    for i in range(L_NS):
        r = R0 + i
        k = K(r)
        site = (f'=IFERROR(IF({KSEL}>0,IF({k}<={KSEL},INDEX($AJ$3:$AJ$10,MATCH({k},$AL$3:$AL$10,0)),""),'
                f'IF({k}<={KPA},INDEX({PJ("AO")},1000-MOD(LARGE({PJ("AU")},{k}),1000)),"")),"")')
        hp(ws, f'AX{r}', site)
        hp(ws, f'AY{r}', f'=IF(AX{r}="","",{esc(f"AX{r}")})')
        hp(ws, f'AZ{r}', f'=IF(AND({KSEL}=0,{k}={KPA}+1),1,0)')
        ws[f'V{r}'] = f'=IF(AX{r}<>"",AX{r},IF(AZ{r}=1,"{CO_NAME}",""))'
        ws[f'W{r}'] = (f'=IF(AX{r}<>"",' + sum8(f'SUMIFS(考_天数{{i}},考_项目{{i}},AY{r},{RNG},考_过账,0)') +
                       f',IF(AZ{r}=1,SUMIFS(考_总天数,{RNG},考_过账,0)-SUM({PJ("AR")}),""))')
        ws[f'X{r}'] = (f'=IF(AX{r}<>"",' + sum8(f'SUMIFS(考_金额{{i}},考_项目{{i}},AY{r},{RNG})') +
                       f',IF(AZ{r}=1,SUMIFS(考_未分摊,{RNG}),""))')
        ws[f'Y{r}'] = (f'=IF(AX{r}<>"",{renci(f"AY{r}", CRIT)},'
                       f'IF(AZ{r}=1,COUNTIFS({RNG},考_未分摊,"<>0"),""))')
        ws[f'Z{r}'] = f'=IF(N(W{r})=0,"",X{r}/W{r})'
    body(ws, R0, s1, 'V', 'Z', 'V', {'W': DAYF, 'X': MONEY, 'Y': INT, 'Z': MONEY}, {'V': AL})

    # ── 左边 ② 按人 × 工地 ──
    # ② 的 8 个工地＝① 的前 8 行；配对用的名单（空位放一个不会重名的记号）
    site = lambda k: f'$AX${R0 + k}'
    sesc = lambda k: f'$AY${R0 + k}'
    for k in range(NPAIR):
        hp(ws, f'BA{R0 + k}', f'=IF({site(k)}="","§",{site(k)})')
    # 考勤逐条：这一行在 ② 这几个工地的天数（第 R0 行起 N_ATT 行）
    for i in range(N_ATT):
        r = R0 + i
        n = K(r)
        hp(ws, f'BD{r}', f'=IF(INDEX(考_有效,{n})=0,0,' +
           sum8(f'IF(ISNUMBER(MATCH(INDEX(考_项目{{i}},{n}),$BA${R0}:$BA${R0 + NPAIR - 1},0)),INDEX(考_天数{{i}},{n}),0)') + ')')
    KQ_D = f'$BD${R0}:$BD${R0 + N_ATT - 1}'
    for i in range(N_PER):
        r = R0 + i
        hp(ws, f'BG{r}', f'=INDEX(人员_姓名,{K(r)})&""')
        hp(ws, f'BH{r}', f'=IF(BG{r}="","",{esc(f"BG{r}")})')
    counter(ws, 'BJ', R0, N_PER, lambda i: f'AND(BG{R0 + i}<>"",SUMIFS({KQ_D},考_姓名,BH{R0 + i},{RNG},考_过账,0)<>0)')
    KN = cnt('BJ', R0, N_PER)
    section(ws, L_SEC, 'A', 'T', '② 按人 × 工地', C_VIEW)
    ws[f'A{L_SEC}'] = (f'="② 按人 × 工地（共 "&{KN}&" 人"&IF({KN}>{L_N},"，只显示前 {L_N} 人","")'
                       f'&IF(AND({KSEL}=0,{KPA}>{NPAIR}),"；没选工地，只列人工成本最大的 {NPAIR} 个工地，看别的请在上面选","")&"）"')
    merge_put(ws, f'A{L_H1}:A{L_H2}', '序号', F_HDR, fill(C_VIEW), align=ACW)
    merge_put(ws, f'B{L_H1}:B{L_H2}', '姓名', F_HDR, fill(C_VIEW), align=ACW)
    for k in range(NPAIR):
        d, a = L_DC[k], L_AC[k]
        merge_put(ws, f'{d}{L_H1}:{a}{L_H1}', f'=IF({site(k)}="","（工地{k + 1}）",{site(k)})', F_HDR, fill(C_VIEW), align=ACW)
        put(ws, f'{d}{L_H2}', '天数', F_HDR, fill(C_VIEW), align=ACW)
        put(ws, f'{a}{L_H2}', '人工', F_HDR, fill(C_VIEW), align=ACW)
    merge_put(ws, f'S{L_H1}:T{L_H1}', '合计（这几个工地）', F_HDR, fill(C_VIEW), align=ACW)
    put(ws, f'S{L_H2}', '天数', F_HDR, fill(C_VIEW), align=ACW)
    put(ws, f'T{L_H2}', '人工', F_HDR, fill(C_VIEW), align=ACW)
    ws.row_dimensions[L_H1].height = 32
    rng = lambda c: f'{c}{R0}:{c}{last}'
    merge_put(ws, f'A{rt}:B{rt}', f'="合计（"&MIN({KN},{L_N})&" 人）"', F_TXTB, FILL_TOT)
    tc = {}
    fm = {}
    for k in range(NPAIR):
        for c, f_ in ((L_DC[k], DAYF), (L_AC[k], MONEY)):
            tc[f'{c}{rt}'] = f'=IF({site(k)}="","",SUM({rng(c)}))'
            fm[c] = f_
    tc[f'S{rt}'], tc[f'T{rt}'] = f'=SUM({rng("S")})', f'=SUM({rng("T")})'
    fm['S'], fm['T'] = DAYF, MONEY
    tot(ws, tc, fm)
    crit_p = lambda r: f'考_姓名,$BO{r},{RNG},'
    for i in range(L_N):
        r = R0 + i
        k = K(r)
        hp(ws, f'BM{r}', '=' + kth(k, 'BJ', R0, N_PER))
        hp(ws, f'BN{r}', f'=IF(BM{r}=0,"",INDEX($BG${R0}:$BG${R0 + N_PER - 1},BM{r}))')
        hp(ws, f'BO{r}', f'=IF(BM{r}=0,"",INDEX($BH${R0}:$BH${R0 + N_PER - 1},BM{r}))')
        ws[f'A{r}'] = f'=IF($BM{r}=0,"",{k})'
        ws[f'B{r}'] = f'=$BN{r}'
        for kk in range(NPAIR):
            blank = f'OR($BM{r}=0,{site(kk)}="")'
            ws[f'{L_DC[kk]}{r}'] = (f'=IF({blank},"",' +
                                    sum8(f'SUMIFS(考_天数{{i}},考_项目{{i}},{sesc(kk)},考_姓名,$BO{r},{RNG},考_过账,0)') + ')')
            ws[f'{L_AC[kk]}{r}'] = (f'=IF({blank},"",' +
                                    sum8(f'SUMIFS(考_金额{{i}},考_项目{{i}},{sesc(kk)},考_姓名,$BO{r},{RNG})') + ')')
        ws[f'S{r}'] = f'=IF($BM{r}=0,"",SUM(' + ','.join(f'{c}{r}' for c in L_DC) + '))'
        ws[f'T{r}'] = f'=IF($BM{r}=0,"",SUM(' + ','.join(f'{c}{r}' for c in L_AC) + '))'
    body(ws, R0, last, 'A', 'T', 'B', fm)
    note(ws, 'A6:T6', f'=IF({KSEL}>0,"实际：选了 "&{KSEL}&" 个工地（按你选的顺序）",'
                      f'"实际：没选工地＝全部工程项目（这段时间 "&{KPA}&" 个工地有用工"'
                      f'&IF({KPA}>{NPAIR},"，② 只列人工成本最大的 {NPAIR} 个","")&"）")'
                      f'&IF(ROUND(T{rt}-SUMIFS(X{R0}:X{R0 + NPAIR - 1},AZ{R0}:AZ{R0 + NPAIR - 1},0),0)<>0,"　⚠ ② 合计跟 ① 前 8 个工地对不上：考勤里有名字不在【人员信息】里","")',
         font=F_AUTOB)

    hide_range(ws, 'AG', 'BZ')
    ws.column_dimensions['BJ'].hidden = True
    ws.freeze_panes = f'C{R0}'
    print_setup(ws, f'{L_H1}:{L_TOT}', landscape=True)
    ws.print_area = f'A1:{L_LAST}{last}'


# ═════════════════════════ 欠薪补发（用户「未发清工资」格式） ═════════════════════════
A_HDR, A_TOT, A_R0, A_N = 5, 6, 7, 40
A_MC = [CL(6 + m) for m in range(12)]        # F..Q：1～12 月补发
A_LAST = 'V'


def _d(v):
    if v in (None, ''):
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    try:
        return dt.datetime.strptime(str(v)[:10], '%Y-%m-%d')
    except ValueError:
        return None


def build_arrear(ws, ctx):
    R0, N = A_R0, A_N
    last = R0 + N - 1
    W = {'A': 5, 'B': 9, 'C': 12, 'D': 11, 'E': 11, 'R': 12, 'S': 12, 'T': 13, 'U': 12, 'V': 30}
    for c in A_MC:
        W[c] = 10
    widths(ws, W)
    tip = ('💡 你那张「未发清工资表」就记在这里：淡黄格填姓名、未发清金额、截至哪天（2025 年底没发清的填 2025/12/31）。'
           '以后补发欠薪不用在这张表记——照常在【收支登记】记一笔工资发放（收支项目用「项目-劳务费」，想单独看也可以在【基本信息】④ 加一个「补发欠薪」、归类选「工资发放」），'
           '人员选他。这张表自动按月算：每月补发＝这个月发给他的－这个月考勤应发（负数＝这个月没发够，又欠多了），剩余＝未发清－补发合计。'
           '右边「系统算的」是按考勤和发放倒算的截至日欠薪，跟你填的对不上就核对一下（看差额）。黄格「补发年度」选看哪一年。')
    title(ws, '欠薪补发', A_LAST, C_VIEW, tip)
    home_link(ws, 'W1')
    ws.column_dimensions['W'].width = 10

    YR, YS = '$AH$3', '$AH$4'
    hp(ws, 'AH3', f'=IF(ISNUMBER(C3),IF(C3>3000,YEAR(C3),INT(C3)),IF(COUNT(D{R0}:D{last})>0,YEAR(MAX(D{R0}:D{last}))+1,P_年度))')
    hp(ws, 'AH4', f'=DATE({YR},1,1)')
    selector(ws, 'B3', '补发年度', 'C3', None)
    dvy = DataValidation(type='whole', operator='between', formula1='2020', formula2='2040', allow_blank=True,
                         showErrorMessage=True, errorStyle='warning', errorTitle='年度', error='填年份，比如 2026')
    ws.add_data_validation(dvy)
    dvy.add('C3')
    note(ws, 'D3:K3', f'="实际："&{YR}&" 年"&IF(ISNUMBER(C3),"","（空着＝截至日期的下一年）")&"；"&{YR}&" 年以后的月份要看就把年度改大"', font=F_AUTOB)
    note(ws, 'A4:V4', '以后补发欠薪：在【收支登记】记工资发放、人员选他，这里自动按月显示；每月补发＝当月已发－当月应发；剩余＝未发清－补发合计；'
                      '首页截止日以后的月份不显示。', align=ALW)
    ws.row_dimensions[4].height = 30

    heads = [('A', '序号'), ('B', '姓名'), ('C', '未发清金额'), ('D', '截至日期'),
             ('E', f'="以前年度补发\n（"&{YR}&"年以前）"')]
    for m in range(12):
        heads.append((A_MC[m], f'={YR}&"年{m + 1}月补发"'))
    heads += [('R', '补发合计'), ('S', '剩余工资\n（未发清－补发合计）'), ('T', '系统算的截至日欠薪\n（考勤应发－已发）'),
              ('U', '差额\n（未发清－系统算的）'), ('V', '说明')]
    header(ws, A_HDR, heads, C_VIEW, height=44)

    rt = A_TOT
    merge_put(ws, f'A{rt}:B{rt}', '合计', F_TXTB, FILL_TOT)
    tc = {f'{c}{rt}': f'=SUM({c}{R0}:{c}{last})' for c in ['C', 'E'] + A_MC + ['R', 'S', 'T', 'U']}
    tc[f'D{rt}'] = None
    fm = {c: MONEY for c in ['C', 'E'] + A_MC + ['R', 'S', 'T', 'U']}
    tot(ws, tc, fm)
    put(ws, f'V{rt}', None, F_TXTB, FILL_TOT)

    for i in range(N):
        r = R0 + i
        nm, e, cut, m0 = f'$AI{r}', f'$AJ{r}', f'$AK{r}', f'$AL{r}'
        hp(ws, f'AI{r}', f'=TRIM(B{r}&"")')
        hp(ws, f'AJ{r}', f'=IF(AI{r}="","",{esc(f"AI{r}")})')
        hp(ws, f'AK{r}', f'=IF(ISNUMBER(D{r}),INT(D{r}),DATE({YR}-1,12,31))')
        hp(ws, f'AL{r}', f'=DATE(YEAR(AK{r}),MONTH(AK{r})+1,1)')
        hp(ws, f'AM{r}', f'=IF(AI{r}="",0,IF(COUNTIF(考_姓名,AJ{r})>0,1,0))')
        ws[f'A{r}'] = f'=IF({nm}="","",COUNTIF($AI${R0}:AI{r},"?*"))'
        ws[f'E{r}'] = (f'=IF(OR({nm}="",{m0}>={YS}),"",ROUND(-SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},收_日期,">="&{m0},收_日期,"<"&{YS})'
                       f'-SUMIFS(考_应发,考_姓名,{e},考_月,">="&{m0},考_月,"<"&{YS}),2))')
        for m in range(12):
            ms = f'DATE({YR},{m + 1},1)'
            ym = f'{YR}*100+{m + 1}'
            ws[f'{A_MC[m]}{r}'] = (f'=IF(OR({nm}="",{ms}<{m0},{ms}>P_截止),"",ROUND(-SUMIFS(收_净额,收_归类,"工资发放",收_人员,{e},收_年月,{ym})'
                                   f'-SUMIFS(考_应发,考_姓名,{e},考_年月,{ym}),2))')
        ws[f'R{r}'] = f'=IF({nm}="","",SUM(E{r}:{A_MC[-1]}{r}))'
        ws[f'S{r}'] = f'=IF({nm}="","",N(C{r})-R{r})'
        ws[f'T{r}'] = f'=IF({nm}="","",ROUND({owed(e, cut, f"(YEAR({cut})*100+MONTH({cut}))", f"$AM{r}")},2))'
        ws[f'U{r}'] = f'=IF({nm}="","",N(C{r})-T{r})'
    # 录入格（淡黄，make 会解锁）＋自动格
    for r in range(R0, last + 1):
        for c in ('B', 'C', 'D', 'V'):
            put(ws, f'{c}{r}', None, F_IN, FILL_IN, {'C': MONEY, 'D': DATE}.get(c), {'V': AL, 'C': AR}.get(c, AC))
        for c in ['A', 'E'] + A_MC + ['R', 'S', 'T', 'U']:
            put(ws, f'{c}{r}', None, F_AUTO, FILL_AUTO, MONEY if c != 'A' else None, AR if c != 'A' else AC)
    dv_list(ws, f'B{R0}:B{last}', '=人员列表', stop=False)
    dv_date(ws, f'D{R0}:D{last}')
    dvm = DataValidation(type='decimal', operator='between', formula1='-999999999', formula2='999999999', allow_blank=True,
                         showErrorMessage=True, errorStyle='stop', errorTitle='金额', error='金额要填数字')
    ws.add_data_validation(dvm)
    dvm.add(f'C{R0}:C{last}')
    # 预填（用户原表）
    rows = ctx.get('arrears') or []
    for i, x in enumerate(rows[:N]):
        r = R0 + i
        ws[f'B{r}'] = x.get('姓名')
        amt = x.get('金额', x.get('未发清金额'))
        if amt not in (None, ''):
            ws[f'C{r}'] = float(amt)
        ws[f'D{r}'] = _d(x.get('截至')) or dt.datetime(2025, 12, 31)
        memo = x.get('说明') or x.get('备注')
        if memo:
            ws[f'V{r}'] = memo
    note(ws, f'A{last + 2}:V{last + 2}', '「系统算的截至日欠薪」＝建账前欠薪（【人员信息】）＋考勤起算月到截至日的考勤应发－同期发给他的工资；'
                                        '跟你填的「未发清金额」差得多，多半是考勤没录全、或者有工资记在别人名下／没填人员。', align=ALW)
    ws.row_dimensions[last + 2].height = 30

    hide_range(ws, 'AH', 'AM')
    ws.freeze_panes = f'C{R0}'
    print_setup(ws, f'{A_HDR}:{A_TOT}', landscape=True)
    ws.print_area = f'A1:{A_LAST}{last + 2}'


def build(wb, ctx=None):
    ctx = ctx or {}
    build_paysum(wb[SH_PAYSUM])
    build_pay(wb[SH_PAY])
    build_labor(wb[SH_LABOR])
    build_arrear(wb[SH_ARREAR], ctx)
