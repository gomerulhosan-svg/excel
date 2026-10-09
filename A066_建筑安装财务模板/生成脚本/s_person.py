# -*- coding: utf-8 -*-
"""查看表【个人往来】（口径：报表口径.md §6）
   ① 汇总：有往来的人一人一行——借款（公司欠他）、垫付报销（他的个人户＋公司还他的）、备用金、过账（代收代付）、往来款、合计公司欠他；
   ② 借款分年（选中的人；没选＝全部人）；③ 选中的人的每一笔明细（按日期排）；顶上「怎么记」。
   只用定义名称 ＋ 本表格子。逐个账户、逐个人、逐条流水的判断放在本表右边隐藏列（BA 列以后）。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *

M = 30                  # ① 最多列几个人
NDET = 500              # ③ 明细最多显示几条
H0 = 10                 # 右边隐藏辅助区的起始行
LAST = 'T'              # 可见区最后一列
COSTK = ['材料', '分包', '机械', '人工', '其他直接费', '管理费用', '财务费用', '营业外支出', '未分类']   # 收_成本类 的全部取值

# ── 行 ──
R_SEL, R_EFF = 3, 4
R_HOW = 5
HOW = [
    '① 老板/员工自己垫钱：在【收支登记】记，账户选他自己的个人户（聂辉微信、王伟微信…；没有就在【基本信息】②加一个，类型「个人户」、所属人选他），'
    '收支项目照常选（材料、分包、招待费…），拿到发票填「已开票金额」「开票日期」。',
    '② 公司还他垫的钱：【收支登记】用公司账户记支出，收支项目选归类「报销还款」的那个，人员选他。他手上替公司收的钱交回公司：记内部转账（账户＝他的个人户，对方账户＝公司账户）。',
    '③ 个人借款：收支项目选「个人借款-XX」（归类「个人借款」）。借进来记收入（或照老习惯在支出栏记负数），还他记支出；人员空着就按收支项目「-」后面的名字。',
    '④ 备用金：领记支出、退记收入，收支项目选归类「备用金」的，人员选他。',
    '⑤ 过账人员（工资、社保只在公司账上过一下的人，比如姚俊强）：【人员信息】「过账人员」选「是」；公司替他付的工资、社保照常记支出，他退回来的记收入（归类「代收代付」），'
    '人员都选他——「过账」一栏自动算还有多少没退。',
]
R_S1 = R_HOW + len(HOW) + 2        # ① 标题条
R_G1, R_H1, R_T1 = R_S1 + 1, R_S1 + 2, R_S1 + 3
R_P0 = R_T1 + 1                    # ① 第 1 个人
R_P1 = R_P0 + M - 1
R_S2 = R_P1 + 2                    # ② 标题条
R_H2 = R_S2 + 1
R_Y0 = R_H2 + 1                    # ② 第 1 年
R_T2 = R_Y0 + NYEARS
R_S3 = R_T2 + 2                    # ③ 标题条
R_H3, R_T3 = R_S3 + 1, R_S3 + 2
R_D0 = R_T3 + 1                    # ③ 第 1 条
R_D1 = R_D0 + NDET - 1

# ── 本表格子 ──
Q, Z = f'$D${R_EFF}', f'$F${R_EFF}'           # 有效起、止
SEL, SELE, DIFF = '$BA$3', '$BA$4', '$BA$5'    # 选中的人、转义后、合计对不上标记
D_PER = f'收_日期,">="&{Q},收_日期,"<="&{Z}'
D_BEF = f'收_日期,"<"&{Q}'
D_UPTO = f'收_日期,"<="&{Z}'
POS, NEG = '收_净额,">0"', '收_净额,"<0"'

# 隐藏：账户块（第 H0 行起 N_ACC 行）
A_NM, A_TY, A_OW, A_OP, A_INB, A_INP, A_INU, A_HIS = 'BA', 'BB', 'BC', 'BD', 'BE', 'BF', 'BG', 'BH'
AR_ = lambda c: f'${c}${H0}:${c}${H0 + N_ACC - 1}'
# 隐藏：人员块（第 H0 行起 N_PER 行）：姓名、计数（counter）、转义
P_NM, P_CNT, P_ESC = 'BJ', 'BK', 'BL'
P_LASTCNT = f'${P_CNT}${H0 + N_PER - 1}'
S_ESC = 'BN'                                   # ① 每个人那行的转义名
K_COL, K_N = 'BQ', 'BR'                        # ③ 排序键（第 H0 行起 N_CASH 行）、第 k 条是 _收 第几条


def fx(pe):
    """① 各列公式（不带等号）。pe＝转义后的人名格；None＝全部人合计"""
    by = lambda rng: f',{rng},{pe}' if pe else ''
    who = by('收_人员')
    own = '收_账户类型,"个人户"' + by('收_账户人')
    hacc = f'{AR_(A_TY)},"个人户"' + by(AR_(A_OW))
    op = lambda t: f'SUMIFS(期初_金额,期初_类型,"{t}"{by("期初_对象")})'
    cat = lambda g, *ex: f'SUMIFS(收_净额,收_归类,"{g}"{who},{",".join(ex)})'
    gz = lambda *ex: f'SUMIFS(收_净额,收_归类,"工资发放",收_过账,1{who},{",".join(ex)})'
    f = {}
    # 借款（公司欠他）
    f['C'] = f'{op("个人借款")}+{cat("个人借款", D_BEF)}'
    f['D'] = cat('个人借款', POS, D_PER)
    f['E'] = '-' + cat('个人借款', NEG, D_PER)
    f['F'] = f'{op("个人借款")}+{cat("个人借款", D_UPTO)}'
    # 垫付报销：个人户余额（含内部转账）＋公司已还他 − 期初公司欠他的垫付；正＝他手上有公司的钱
    bal = lambda d, hc: (f'SUMIFS({AR_(A_OP)},{hacc})+SUMIFS(收_净额,{own},{d})+SUMIFS({AR_(hc)},{hacc})'
                         f'-{cat("报销还款", d)}-{op("个人垫付")}')
    f['G'] = bal(D_BEF, A_INB)
    f['H'] = f'SUMIFS(收_支出,{own},收_归类,"<>内部转账",{D_PER})'
    f['I'] = f'SUMIFS(收_收入,{own},收_归类,"<>内部转账",{D_PER})'
    f['J'] = f'SUMIFS(收_净额,{own},收_归类,"内部转账",{D_PER})+SUMIFS({AR_(A_INP)},{hacc})'
    f['K'] = '-' + cat('报销还款', D_PER)
    f['L'] = bal(D_UPTO, A_INU)
    f['M'] = f'SUMIFS(收_已开票,{own},收_支出,"<>0",{D_PER})'
    f['N'] = ('MAX(0,-(' + '+'.join(f'SUMIFS(收_净额,{own},收_成本类,"{k}",{D_PER})+SUMIFS(收_已开票,{own},收_成本类,"{k}",{D_PER})'
                                    for k in COSTK) + '))')
    # 备用金（他手上的）
    f['O'] = f'{op("备用金")}-{cat("备用金", D_UPTO)}'
    # 过账（代收代付 ＋ 过账人员的工资发放）
    f['P'] = f'-{cat("代收代付", NEG, D_PER)}-{gz(NEG, D_PER)}'
    f['Q'] = f'{cat("代收代付", POS, D_PER)}+{gz(POS, D_PER)}'
    f['R'] = f'{op("代收代付")}-{cat("代收代付", D_UPTO)}-{gz(D_UPTO)}'
    # 往来款（正＝他欠公司）
    f['S'] = f'{op("其他应收")}-{op("其他应付")}-SUMIFS(收_净额,收_归类,"往来款"{by("收_往来对象")},{D_UPTO})'
    return f


def build(wb, ctx=None):
    ws = wb[SH_PERSON]
    W = {'A': 11, 'B': 12, 'C': 13, 'D': 12.5, 'E': 12.5, 'F': 12.5, 'G': 12.5, 'H': 13, 'I': 15, 'J': 13, 'K': 12.5,
         'L': 13, 'M': 12, 'N': 12, 'O': 12, 'P': 12, 'Q': 12, 'R': 12, 'S': 12.5, 'T': 13.5, 'U': 10}
    widths(ws, W)
    tip = ('💡 每个跟公司有钱来往的人一行：借款（公司欠他多少）、垫付报销（他自己的微信/银行卡替公司付的、收的，公司已经还他的，开没开票）、'
           '备用金、过账（代收代付）、往来款，最后一列合计公司欠他多少。黄格：选「人员」，下面就列出他的借款分年和每一笔明细；起止空着＝建账日到首页截止日。')
    title(ws, '个人往来', LAST, C_VIEW, tip)
    home_link(ws, 'U1')

    # ── 选择格 ──
    selector(ws, f'A{R_SEL}', '人员', f'B{R_SEL}', None, dv_formula='=人员列表',
             prompt='选一个人：下面列他的借款分年和每一笔明细；空着＝不看明细（借款分年显示全部人合计）')
    selector(ws, f'C{R_SEL}', '起', f'D{R_SEL}', None, fmt=DATE)
    selector(ws, f'E{R_SEL}', '止', f'F{R_SEL}', None, fmt=DATE)
    dv_date(ws, f'D{R_SEL}')
    dv_date(ws, f'F{R_SEL}')
    put(ws, f'A{R_EFF}', '实际用的', F_NOTE, align=AC)
    put(ws, f'B{R_EFF}', f'=IF({SEL}="","（空＝不看明细）",{SEL})', F_AUTOB, FILL_AUTO, align=AC)
    put(ws, f'C{R_EFF}', '起', F_NOTE, align=AC)
    put(ws, f'D{R_EFF}', f'=IF(ISNUMBER(D{R_SEL}),INT(D{R_SEL}),P_建账日)', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, f'E{R_EFF}', '止', F_NOTE, align=AC)
    put(ws, f'F{R_EFF}', f'=IF(ISNUMBER(F{R_SEL}),INT(F{R_SEL}),P_截止)', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells(f'G{R_SEL}:{LAST}{R_EFF}')
    put(ws, f'G{R_SEL}', '起空着＝建账日（【基本信息】①），止空着＝首页截止日。① 汇总和 ③ 明细按这段日期算（期初＝起那天以前的余额，期末＝止那天的余额）；② 借款分年按年算到止那天。',
        F_NOTE, align=ALW, border=False)
    ws[SEL] = f'=TRIM(B{R_SEL}&"")'
    ws[SELE] = f'={esc(SEL)}'

    # ── 怎么记 ──
    section(ws, R_HOW, 'A', LAST, '📌 怎么记（老板垫付、公司还钱、借款、备用金、过账，各一句话）', C_VIEW)
    for i, t in enumerate(HOW):
        r = R_HOW + 1 + i
        ws.merge_cells(f'A{r}:{LAST}{r}')
        put(ws, f'A{r}', t, F_TXT, FILL_TIP, align=ALW)
        ws.row_dimensions[r].height = 30 if len(t) > 110 else 18

    _hidden_accounts(ws)
    _hidden_persons(ws)
    _summary(ws)
    _loans_by_year(ws)
    _detail(ws)

    hide(ws, *[CL(i) for i in range(CI('BA'), CI('CA') + 1)])
    ws.freeze_panes = f'C{R_EFF + 1}'
    print_setup(ws, f'1:{R_EFF}', landscape=True)


# ───────────────────────── 隐藏：账户块、人员块 ─────────────────────────
def _hidden_accounts(ws):
    """每个账户：名称、类型、所属人、期初；别的账户「内部转账」转进它的钱（起前、本期、到止）；是不是选中的人的个人户"""
    for c, t in zip((A_NM, A_TY, A_OW, A_OP, A_INB, A_INP, A_INU, A_HIS),
                    ('账户', '类型', '所属人', '期初', '转进_起前', '转进_本期', '转进_到止', '选中人的个人户')):
        ws[f'{c}{H0 - 1}'] = t
        ws[f'{c}{H0 - 1}'].font = F_HELP
    for j in range(N_ACC):
        r = H0 + j
        a = f'{A_NM}{r}'
        ws[a] = f'=INDEX(账户_名称,{j + 1})&""'
        ws[f'{A_TY}{r}'] = f'=INDEX(账户_类型,{j + 1})&""'
        ws[f'{A_OW}{r}'] = f'=INDEX(账户_所属人,{j + 1})&""'
        ws[f'{A_OP}{r}'] = f'=IF({a}="",0,N(INDEX(账户_期初,{j + 1})))'
        for c, d in ((A_INB, D_BEF), (A_INP, D_PER), (A_INU, D_UPTO)):
            ws[f'{c}{r}'] = f'=IF({A_TY}{r}<>"个人户",0,-SUMIFS(收_净额,收_对方账户,{esc(a)},收_归类,"内部转账",{d}))'
        ws[f'{A_HIS}{r}'] = f'=IF(AND({SEL}<>"",{A_TY}{r}="个人户",{A_OW}{r}={SEL}),{a},"")'
        for c in (A_NM, A_TY, A_OW, A_OP, A_INB, A_INP, A_INU, A_HIS):
            ws[f'{c}{r}'].font = F_HELP


def _hidden_persons(ws):
    """【人员信息】每个人：到止那天为止有没有任何往来（借款、报销还款、备用金、代收代付、过账工资、往来款、个人户、期初）"""
    ws[f'{P_NM}{H0 - 1}'], ws[f'{P_CNT}{H0 - 1}'], ws[f'{P_ESC}{H0 - 1}'] = '人员', '有往来的第几个', '转义'
    for i in range(N_PER):
        r = H0 + i
        ws[f'{P_NM}{r}'] = f'=INDEX(人员_姓名,{i + 1})&""'
        ws[f'{P_ESC}{r}'] = f'={esc(f"{P_NM}{r}")}'
        ws[f'{P_NM}{r}'].font = ws[f'{P_ESC}{r}'].font = F_HELP

    def cond(i):
        r = H0 + i
        e = f'{P_ESC}{r}'
        parts = [f'COUNTIFS(收_人员,{e},收_归类,"{g}",{D_UPTO})' for g in ('个人借款', '报销还款', '备用金', '代收代付')]
        parts.append(f'COUNTIFS(收_人员,{e},收_归类,"工资发放",收_过账,1,{D_UPTO})')
        parts.append(f'COUNTIFS(收_往来对象,{e},收_归类,"往来款",{D_UPTO})')
        parts.append(f'COUNTIFS({AR_(A_OW)},{e},{AR_(A_TY)},"个人户")')
        parts += [f'COUNTIFS(期初_对象,{e},期初_类型,"{t}")' for t in ('个人借款', '个人垫付', '备用金', '代收代付', '其他应收', '其他应付')]
        return f'IF({P_NM}{r}="",FALSE,{"+".join(parts)}>0)'

    counter(ws, P_CNT, H0, N_PER, cond)


# ───────────────────────── ① 汇总 ─────────────────────────
GROUPS1 = [('A', 'B', '人'), ('C', 'F', '借款（公司欠他的）'),
           ('G', 'N', '垫付报销（他的个人户＋公司还他的；余额 负＝公司还欠他，正＝他手上有公司的钱）'),
           ('O', 'O', '备用金'), ('P', 'R', '过账（代收代付）'), ('S', 'S', '往来款'), ('T', 'T', '合计')]
HEAD1 = [('A', '姓名'), ('B', '类别'),
         ('C', '期初欠他'), ('D', '本期借入'), ('E', '本期归还'), ('F', '期末欠他'),
         ('G', '期初余额'), ('H', '本期他垫付\n（个人户付出去）'), ('I', '本期替公司收\n（进他个人户）'), ('J', '本期内部转账\n（转进他户为正）'),
         ('K', '本期公司\n已还他'), ('L', '期末余额\n＝期初−垫付＋收\n±转账＋已还'), ('M', '已开票\n（垫付里）'), ('N', '欠票\n（费用没票的）'),
         ('O', '期末他手上\n（他欠公司）'),
         ('P', '本期公司\n替他付'), ('Q', '本期他退回'), ('R', '未退\n（他欠公司）'),
         ('S', '余额\n（正＝他欠公司）'),
         ('T', '公司欠他\n（负＝他欠公司）')]
MONEY_COLS1 = [c for c, _ in HEAD1[2:]]
DIRECT1 = list('CDEFGHIJKLMOPQR')          # 合计行直接按全部算的列（其余＝各人加起来）


def _summary(ws):
    cnt = P_LASTCNT
    rng = lambda c: f'{c}{R_P0}:{c}{R_P1}'
    diffs = ','.join(f'ROUND({c}{R_T1}-SUM({rng(c)}),2)<>0' for c in DIRECT1)
    ws[DIFF] = f'=IF(OR({diffs}),1,0)'
    section(ws, R_S1, 'A', LAST,
            f'="① 汇总："&TEXT({Q},"yyyy-mm-dd")&" ～ "&TEXT({Z},"yyyy-mm-dd")&"，有往来的 "&{cnt}&" 人（工资不在这里，看【工资表】）"'
            f'&IF({cnt}>{M},"　⚠ 超过 {M} 人，只列了前 {M} 个","")'
            f'&IF({DIFF}=1,"　⚠ 合计里有钱没对上人：人员没填或不在【人员信息】里（看【收支登记】校验）","")', C_VIEW)
    for c1, c2, t in GROUPS1:
        put(ws, f'{c1}{R_G1}', t, F_HDR, fill(C_VIEW), align=ACW)
        for i in range(CI(c1), CI(c2) + 1):
            ws.cell(row=R_G1, column=i).border = BD
            ws.cell(row=R_G1, column=i).fill = fill(C_VIEW)
        if c1 != c2:
            ws.merge_cells(f'{c1}{R_G1}:{c2}{R_G1}')
    ws.row_dimensions[R_G1].height = 22
    header(ws, R_H1, HEAD1, 'FF70AD47', height=46)

    # 合计行
    tot = fx(None)
    put(ws, f'A{R_T1}', '合计', F_AUTOB, FILL_TOT, align=AC)
    put(ws, f'B{R_T1}', f'={cnt}&" 人"', F_AUTOB, FILL_TOT, align=AC)
    for c in MONEY_COLS1:
        if c in DIRECT1:
            v = f'=ROUND({tot[c]},2)'
        elif c == 'T':
            v = f'=ROUND(F{R_T1}-L{R_T1}-O{R_T1}-R{R_T1}-S{R_T1},2)'
        else:
            v = f'=ROUND(SUM({rng(c)}),2)'
        put(ws, f'{c}{R_T1}', v, F_AUTOB, FILL_TOT, MONEY, AR)

    # 每个人
    for k in range(M):
        r = R_P0 + k
        a = f'$A{r}'
        e = f'${S_ESC}{r}'
        ws[f'A{r}'] = (f'=IF({k + 1}>{cnt},"",IFERROR(INDEX(${P_NM}${H0}:${P_NM}${H0 + N_PER - 1},'
                       f'MATCH({k + 1},${P_CNT}${H0}:${P_CNT}${H0 + N_PER - 1},0)),""))')
        ws[f'B{r}'] = f'=IF({a}="","",IFERROR(INDEX(人员_类别,MATCH({a},人员_姓名,0))&"",""))'
        ws[f'{S_ESC}{r}'] = f'=IF({a}="","",{esc(a)})'
        ws[f'{S_ESC}{r}'].font = F_HELP
        f = fx(e)
        for c in MONEY_COLS1:
            if c == 'T':
                v = f'=IF({a}="","",ROUND(F{r}-L{r}-O{r}-R{r}-S{r},2))'
            else:
                v = f'=IF({a}="","",ROUND({f[c]},2))'
            ws[f'{c}{r}'] = v
            ws[f'{c}{r}'].number_format = MONEY
            ws[f'{c}{r}'].font = F_AUTOB if c in ('F', 'L', 'T') else F_AUTO
            ws[f'{c}{r}'].alignment = AR
        ws[f'A{r}'].font = F_TXTB
        ws[f'B{r}'].font = F_AUTO
        ws[f'A{r}'].alignment = ws[f'B{r}'].alignment = AC
    area = f'A{R_P0}:{LAST}{R_P1}'
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'AND($A{R_P0}<>"",$A{R_P0}={SEL})'], fill=fill('FFFFF2CC'), border=BD))
    ws.conditional_formatting.add(area, FormulaRule(formula=[f'$A{R_P0}<>""'], border=BD))


# ───────────────────────── ② 借款分年 ─────────────────────────
def _loans_by_year(ws):
    section(ws, R_S2, 'A', 'E', f'="② 借款分年："&IF({SEL}="","全部人合计",{SEL})&"（公司欠他的）"', C_VIEW)
    header(ws, R_H2, [('A', '年度'), ('B', '年初欠他'), ('C', '借入'), ('D', '归还'), ('E', '年末欠他\n（到止那天）')], 'FF70AD47', height=34)
    ws.merge_cells(f'G{R_Y0}:N{R_Y0 + 2}')
    put(ws, f'G{R_Y0}', '借入＝「个人借款」收进来的钱（支出栏记负数也算借入）；归还＝还他的钱。从建账那年列到「止」那年，最后一年算到止那天。'
                        '选了人就是他一个人的，没选就是全部人合计。', F_NOTE, align=ALW, border=False)
    fy = f'MAX({YEAR0},YEAR(P_建账日))'
    op_all = 'SUMIFS(期初_金额,期初_类型,"个人借款")'
    op_one = f'SUMIFS(期初_金额,期初_类型,"个人借款",期初_对象,{SELE})'
    loan = lambda extra, one: f'SUMIFS(收_净额,收_归类,"个人借款"{f",收_人员,{SELE}" if one else ""},{extra})'
    for i in range(NYEARS):
        r = R_Y0 + i
        y = f'$A{r}'
        ws[f'A{r}'] = f'=IF(AND({fy}+{i}<=YEAR({Z}),{fy}+{i}<={YEAR0 + NYEARS - 1}),{fy}+{i},"")'
        d0 = f'收_日期,"<"&DATE({y},1,1)'
        dy = f'收_日期,">="&DATE({y},1,1),收_日期,"<="&MIN(DATE({y},12,31),{Z})'
        ws[f'B{r}'] = f'=IF({y}="","",ROUND(IF({SEL}="",{op_all}+{loan(d0, 0)},{op_one}+{loan(d0, 1)}),2))'
        dp, dn = f'{dy},{POS}', f'{dy},{NEG}'
        ws[f'C{r}'] = f'=IF({y}="","",ROUND(IF({SEL}="",{loan(dp, 0)},{loan(dp, 1)}),2))'
        ws[f'D{r}'] = f'=IF({y}="","",ROUND(-IF({SEL}="",{loan(dn, 0)},{loan(dn, 1)}),2))'
        ws[f'E{r}'] = f'=IF({y}="","",ROUND(B{r}+C{r}-D{r},2))'
        for c in 'ABCDE':
            ws[f'{c}{r}'].font = F_AUTOB if c in 'AE' else F_AUTO
            ws[f'{c}{r}'].alignment = AC if c == 'A' else AR
            if c != 'A':
                ws[f'{c}{r}'].number_format = MONEY
    ws.conditional_formatting.add(f'A{R_Y0}:E{R_Y0 + NYEARS - 1}', FormulaRule(formula=[f'$A{R_Y0}<>""'], border=BD))
    r = R_T2
    put(ws, f'A{r}', '合计', F_AUTOB, FILL_TOT, align=AC)
    put(ws, f'B{r}', f'=IF(A{R_Y0}="","",B{R_Y0})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'C{r}', f'=SUM(C{R_Y0}:C{R_Y0 + NYEARS - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'D{r}', f'=SUM(D{R_Y0}:D{R_Y0 + NYEARS - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'E{r}', f'=IF(A{R_Y0}="","",ROUND(N(B{r})+C{r}-D{r},2))', F_AUTOB, FILL_TOT, MONEY, AR)


# ───────────────────────── ③ 明细 ─────────────────────────
def _detail(ws):
    his = AR_(A_HIS)
    ix = lambda f, n: f'INDEX(收_{f},{n})'

    def cond(i):
        n = i + 1
        return (f'IF({SEL}="",FALSE,AND({ix("日期", n)}>={Q},{ix("日期", n)}<={Z},OR({ix("人员", n)}={SEL},'
                f'AND({ix("账户类型", n)}="个人户",{ix("账户人", n)}={SEL}),'
                f'AND({ix("归类", n)}="往来款",{ix("往来对象", n)}={SEL}),'
                f'AND({ix("归类", n)}="内部转账",ISNUMBER(MATCH({ix("对方账户", n)},{his},0))))))')

    ncell = skey(ws, K_COL, H0, N_CASH, cond, lambda i: ix('日期', i + 1))
    section(ws, R_S3, 'A', LAST,
            f'=IF({SEL}="","③ 明细：在上面黄格「人员」选一个人，这里就列出他的每一笔（借款、垫付、报销还款、备用金、过账、工资、往来）",'
            f'"③ "&{SEL}&" 的每一笔："&TEXT({Q},"yyyy-mm-dd")&" ～ "&TEXT({Z},"yyyy-mm-dd")&"，共 "&{ncell}&" 条"'
            f'&IF({ncell}>{NDET},"，只显示前 {NDET} 条，请把起止日期缩短","")&"（他的个人户上的每一笔、人员是他的每一笔，按日期排）")', C_VIEW)
    heads = [('A', '日期'), ('B', '账户'), ('C', '收支项目'), ('D', '摘要'), ('E', ''), ('F', ''), ('G', '收入'), ('H', '支出'),
             ('I', '分到哪块'), ('J', '已开票'), ('K', '开票日期'), ('L', '项目'), ('M', '收支登记\n第几行'), ('N', '校验')]
    header(ws, R_H3, heads, 'FF70AD47', height=34)
    ws.merge_cells(f'D{R_H3}:F{R_H3}')
    r = R_T3
    put(ws, f'A{r}', '合计', F_AUTOB, FILL_TOT, align=AC)
    for c in 'BCDEFIKLMN':
        put(ws, f'{c}{r}', None, F_AUTOB, FILL_TOT)
    ws.merge_cells(f'D{r}:F{r}')
    for c in 'GHJ':
        put(ws, f'{c}{r}', f'=SUM({c}{R_D0}:{c}{R_D1})', F_AUTOB, FILL_TOT, MONEY, AR)

    for k in range(NDET):
        r = R_D0 + k
        n = f'${K_N}{r}'
        ws[f'{K_N}{r}'] = f'={ksorted(k + 1, K_COL, H0, N_CASH, ncell)}'
        ws[f'{K_N}{r}'].font = F_HELP
        g = ix('归类', n)
        net = ix('净额', n)
        mine = f'AND({ix("账户类型", n)}="个人户",{ix("账户人", n)}={SEL})'
        cp = f'AND({g}="内部转账",ISNUMBER(MATCH({ix("对方账户", n)},{his},0)))'
        where = (f'IF({ix("有效", n)}=0,"✗ 没算进去（看校验）",'
                 f'IF({mine},IF({g}="内部转账",IF({net}<0,"内部转账·转出","内部转账·转进"),IF({net}<0,"垫付（他的户付出）","替公司收（进他的户）")),'
                 f'IF({cp},IF({net}<0,"内部转账·转进","内部转账·转出"),'
                 f'IF({g}="个人借款",IF({net}>0,"借款·借入","借款·归还"),'
                 f'IF({g}="报销还款","公司还他垫付",'
                 f'IF({g}="备用金",IF({net}<0,"备用金·领","备用金·退"),'
                 f'IF(OR({g}="代收代付",AND({g}="工资发放",{ix("过账", n)}=1)),IF({net}<0,"过账·公司替他付","过账·他退回"),'
                 f'IF({g}="工资发放","工资（不算往来）",IF({g}="往来款","往来款","其他")))))))))')
        vals = {
            'A': (ix('日期', n), DATE, AC), 'B': (ix('账户', n), None, AC), 'C': (ix('收支项目', n), None, AL),
            'D': (ix('摘要', n), None, AL), 'G': (ix('收入', n), MONEY, AR), 'H': (ix('支出', n), MONEY, AR),
            'I': (where, None, AC), 'J': (ix('已开票', n), MONEY, AR),
            'K': (f'IF({ix("开票日期", n)}>0,{ix("开票日期", n)},"")', DATE, AC), 'L': (ix('项目原', n), None, AC),
            'M': (ix('录入行', n), '0', AC), 'N': (ix('校验', n), None, AL)}
        for c, (v, fm, al) in vals.items():
            ws[f'{c}{r}'] = f'=IF({n}=0,"",{v})'
            ws[f'{c}{r}'].font = F_AUTO
            ws[f'{c}{r}'].alignment = al
            if fm:
                ws[f'{c}{r}'].number_format = fm
        ws[f'N{r}'].font = F_RED
        ws.merge_cells(f'D{r}:F{r}')
    ws.conditional_formatting.add(f'A{R_D0}:N{R_D1}', FormulaRule(formula=[f'$A{R_D0}<>""'], border=BD))
