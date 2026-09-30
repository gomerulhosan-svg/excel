# -*- coding: utf-8 -*-
"""录入用的几张登记表：【应付登记】【收入确认】【代发抵账】【发票登记】"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *



def _chk_cf(ws, col, r0, r1):
    rng_ = f'{col}{r0}:{col}{r1}'
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${col}{r0},1)="✗"'], fill=FILL_WARN,
                                                    font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'LEFT(${col}{r0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng_, FormulaRule(formula=[f'${col}{r0}="√"'],
                                                    font=Font(name=YH, sz=10, bold=True, color='FF00B050')))


def _ym(d):
    return f'=IF(ISNUMBER({d}),YEAR({d})*100+MONTH({d}),"")'


def build_ap(wb, ctx):
    ws = wb.create_sheet(SH_AP)
    widths(ws, {'A': 6, 'B': 12, 'C': 10, 'D': 18, 'E': 14, 'F': 30, 'G': 13, 'H': 7, 'I': 22, 'J': 8, 'K': 30})
    title(ws, '应 付 登 记（材料 · 分包 · 机械运输 · 其他直接费：欠人家的钱先记在这里）', AP_CHK, C_AR,
          '💡 跟材料商、分包、机械班组对了账（或者收到送货单、结算单）就登记一行：发生日期、类型、单位、项目、内容、应付金额（含税）。'
          '日期填「发生」的日期（到货、完工或对账所属月份）：12 月的活就算次年才对账、付款，也按 12 月的日期登记；跨几个月的机械费按月拆行。'
          '付款不用在这里记：在【资金流水】付给这家的钱会自动冲它的应付。登记过的单位就是「挂账单位」；'
          '零星现买现付、没登记过的单位，流水里付的钱直接算项目成本。建账日期之前的应付不在这里录，在【期初余额】② 按单位×项目填累计数。')
    header(ws, AP_HDR, [(AP_SEQ, '序号'), (AP_DATE, '发生日期'), (AP_TYPE, '类型'), (AP_UNIT, '单位'), (AP_PJ, '项目'),
                        (AP_MEMO, '内容（材料名称 / 分包内容 / 台班）'), (AP_AMT, '应付金额\n（含税）'), (AP_RATE, '票的税率\n（备查）'),
                        (AP_NOTE, '备注'), (AP_YM, '年月'), (AP_CHK, '校验')], C_AR)
    for r in range(AP_R0, AP_R1 + 1):
        g = lambda c: f'{c}{r}'
        blank = f'AND({g(AP_DATE)}="",{g(AP_UNIT)}="",{g(AP_AMT)}="")'
        ws[g(AP_SEQ)] = f'=IF({blank},"",ROW()-{AP_HDR})'
        ws[g(AP_YM)] = _ym(g(AP_DATE))
        ws[g(AP_CHK)] = (f'=IF({blank},"",IF(NOT(ISNUMBER({g(AP_DATE)})),"✗ 日期要填成日期",'
                         f'IF({g(AP_DATE)}<{OPEN_DATE},"✗ 早于建账日期（建账前的放【期初余额】②）",IF({g(AP_TYPE)}="","✗ 选类型",IF({g(AP_UNIT)}="","✗ 填单位",'
                         f'IF(COUNTIF({UN_NAMES},{g(AP_UNIT)})=0,"✗ 单位没在【往来单位】登记",'
                         f'IF({g(AP_PJ)}="","✗ 选项目（分不清的挑主要那个项目）",'
                         f'IF(COUNTIF({PJ_NAMES},{g(AP_PJ)})=0,"✗ 项目不在【项目档案】里",'
                         f'IF(N({g(AP_AMT)})=0,"✗ 没有金额",'
                         f'"√")))))))))')
    style_rows(ws, AP_R0, AP_R1, [AP_SEQ, AP_DATE, AP_TYPE, AP_UNIT, AP_PJ, AP_MEMO, AP_AMT, AP_RATE, AP_NOTE, AP_YM, AP_CHK],
               auto=[AP_SEQ, AP_YM, AP_CHK], fmts={AP_DATE: DATE, AP_AMT: MONEY, AP_RATE: '0%'},
               aligns={AP_UNIT: AL, AP_MEMO: AL, AP_NOTE: AL, AP_CHK: AL, AP_AMT: AR},
               fills={c: FILL_IN for c in (AP_DATE, AP_TYPE, AP_UNIT, AP_PJ, AP_MEMO, AP_AMT, AP_RATE, AP_NOTE)})
    for r in range(AP_R0, AP_R1 + 1):
        ws[f'{AP_YM}{r}'].font = F_HELP
        ws[f'{AP_UTYPE}{r}'] = f'=IF({AP_UNIT}{r}="","",IFERROR(INDEX({UN_TYPES_R},MATCH({AP_UNIT}{r},{UN_NAMES},0))&"",""))'
        ws[f'{AP_UTYPE}{r}'].font = F_HELP
    ws[f'{AP_UTYPE}{AP_HDR}'] = '单位类型'
    hide(ws, AP_UTYPE)
    _chk_cf(ws, AP_CHK, AP_R0, AP_R1)
    dv_date(ws, f'{AP_DATE}{AP_R0}:{AP_DATE}{AP_R1}')
    dv_list(ws, f'{AP_TYPE}{AP_R0}:{AP_TYPE}{AP_R1}', '"' + ','.join(AP_TYPES) + '"')
    dv_list(ws, f'{AP_UNIT}{AP_R0}:{AP_UNIT}{AP_R1}', f'={UN_NAMES}', '【往来单位】登记过的', stop=False)
    dv_list(ws, f'{AP_PJ}{AP_R0}:{AP_PJ}{AP_R1}', f'={PJ_NAMES}')
    ws.auto_filter.ref = f'A{AP_HDR}:{AP_CHK}{AP_R1}'
    ws.freeze_panes = f'D{AP_R0}'
    for i, row in enumerate(ctx['ap_rows']):
        r = AP_R0 + i
        for c, v in zip((AP_DATE, AP_TYPE, AP_UNIT, AP_PJ, AP_MEMO, AP_AMT, AP_RATE, AP_NOTE), row):
            if v not in (None, ''):
                ws[f'{c}{r}'] = v
    return ws


def build_rev(wb, ctx):
    ws = wb.create_sheet(SH_REV)
    widths(ws, {'A': 6, 'B': 12, 'C': 14, 'D': 11, 'E': 14, 'F': 36, 'G': 26, 'H': 8, 'I': 30})
    title(ws, '收 入 确 认（甲方 / 总包确认的产值、结算 —— 收入按这里算，不按收款）', RV_CHK, C_RPT,
          '💡 甲方或总包确认了进度产值（报量、进度款申请批下来）、签证变更，就登记一行：日期、项目、类型、确认金额（含税）。'
          '利润表、项目利润都按这里的金额算收入；应收账款＝这里确认的 − 收到的工程款。钱几年收不到也不影响收入和利润。'
          '竣工结算了：先在【项目档案】填结算额，再在这里录一行「结算调整」＝结算额 − 以前确认过的合计（多退少补，可以是负数），千万别把结算总额整笔再录一遍。'
          '建账日期之前确认过的产值不在这里录，在【期初余额】① 填累计数。')
    header(ws, RV_HDR, [(RV_SEQ, '序号'), (RV_DATE, '确认日期'), (RV_PJ, '项目'), (RV_TYPE, '类型'), (RV_AMT, '确认金额\n（含税）'),
                        (RV_NOTE, '依据 / 说明'), (RV_CUS, '甲方/总包（自动）'), (RV_YM, '年月'), (RV_CHK, '校验')], C_RPT)
    for r in range(RV_R0, RV_R1 + 1):
        g = lambda c: f'{c}{r}'
        blank = f'AND({g(RV_DATE)}="",{g(RV_PJ)}="",{g(RV_AMT)}="")'
        ws[g(RV_SEQ)] = f'=IF({blank},"",ROW()-{RV_HDR})'
        ws[g(RV_CUS)] = f'=IF({g(RV_PJ)}="","",IFERROR(INDEX({rng(SH_PROJ, PJ_CUS, PJ_R0, PJ_R1)},MATCH({g(RV_PJ)},{PJ_NAMES},0))&"",""))'
        ws[g(RV_YM)] = _ym(g(RV_DATE))
        prate = f'IFERROR(INDEX({rng(SH_PROJ, PJ_TAX, PJ_R0, PJ_R1)},MATCH({g(RV_PJ)},{PJ_NAMES},0)),"")'
        rate = f'IF(ISNUMBER({prate}),{prate},{PA_VATR})'
        tot = f'IFERROR(INDEX({rng(SH_PROJ, PJ_TOTAL, PJ_R0, PJ_R1)},MATCH({g(RV_PJ)},{PJ_NAMES},0)),0)'
        ws[g(RV_VAT)] = f'=IF(N({g(RV_AMT)})=0,0,ROUND({g(RV_AMT)}/(1+{rate})*{rate},2))'
        ws[g(RV_VAT)].font = F_HELP
        ws[g(RV_CHK)] = (f'=IF({blank},"",IF(NOT(ISNUMBER({g(RV_DATE)})),"✗ 日期要填成日期",'
                         f'IF({g(RV_DATE)}<{OPEN_DATE},"✗ 早于建账日期（建账前的放【期初余额】①）",IF({g(RV_PJ)}="","✗ 选项目",IF(COUNTIF({PJ_NAMES},{g(RV_PJ)})=0,"✗ 项目不在【项目档案】里",'
                         f'IF(N({g(RV_AMT)})=0,"✗ 没有金额",IF({g(RV_TYPE)}="","⚠ 选一下类型",'
                         f'IF(AND(N({tot})>0,SUMIFS({rvr(RV_AMT)},{rvr(RV_PJ)},{g(RV_PJ)})+SUMIFS({opr(OPJ_REV)},{opr(OPJ_PJ)},{g(RV_PJ)})>N({tot})*1.001),'
                         f'"⚠ 这个项目累计确认超过总价了（结算额整笔又录了一遍？）","√"))))))))')
    style_rows(ws, RV_R0, RV_R1, [RV_SEQ, RV_DATE, RV_PJ, RV_TYPE, RV_AMT, RV_NOTE, RV_CUS, RV_YM, RV_CHK],
               auto=[RV_SEQ, RV_CUS, RV_YM, RV_CHK], fmts={RV_DATE: DATE, RV_AMT: MONEY},
               aligns={RV_NOTE: AL, RV_CUS: AL, RV_CHK: AL, RV_AMT: AR},
               fills={c: FILL_IN for c in (RV_DATE, RV_PJ, RV_TYPE, RV_AMT, RV_NOTE)})
    for r in range(RV_R0, RV_R1 + 1):
        ws[f'{RV_YM}{r}'].font = F_HELP
    ws[f'{RV_VAT}{RV_HDR}'] = '销项税额'
    ws[f'{RV_VAT}{RV_HDR}'].font = F_HELP
    hide(ws, RV_VAT)
    _chk_cf(ws, RV_CHK, RV_R0, RV_R1)
    dv_date(ws, f'{RV_DATE}{RV_R0}:{RV_DATE}{RV_R1}')
    dv_list(ws, f'{RV_PJ}{RV_R0}:{RV_PJ}{RV_R1}', f'={PJ_NAMES}')
    dv_list(ws, f'{RV_TYPE}{RV_R0}:{RV_TYPE}{RV_R1}', '"' + ','.join(RV_TYPES) + '"')
    ws.auto_filter.ref = f'A{RV_HDR}:{RV_CHK}{RV_R1}'
    ws.freeze_panes = f'D{RV_R0}'
    for i, row in enumerate(ctx['rev_rows']):
        r = RV_R0 + i
        for c, v in zip((RV_DATE, RV_PJ, RV_TYPE, RV_AMT, RV_NOTE), row):
            if v not in (None, ''):
                ws[f'{c}{r}'] = v
    return ws


def build_off(wb, ctx):
    ws = wb.create_sheet(SH_OFF)
    widths(ws, {'A': 6, 'B': 12, 'C': 14, 'D': 18, 'E': 14, 'F': 13, 'G': 36, 'H': 8, 'I': 10, 'J': 30})
    title(ws, '代 发 抵 账（总包代发工资 · 总包代付款 · 甲方扣款：没经过我们账户的钱）', OF_CHK, C_AR,
          '💡 总包直接把工资发给我们的工人（从我们的工程款里扣），选「总包代发工资」：这个项目算收到了这么多工程款，这个工人算已经发了这么多工资。'
          '工人多拿了、退给我们的，在【资金流水】记收入、类别选「工资退回」、单位选这个工人。'
          '总包替我们付给材料商、分包的，选「总包代付材料分包款」（对象填那家单位）；甲方供的材料从工程款里扣，选「甲供材扣款」（算这个项目的材料费）；'
          '甲方扣的罚款、水电费，选「甲方扣款」（算这个项目的其他直接费）。')
    header(ws, OF_HDR, [(OF_SEQ, '序号'), (OF_DATE, '日期'), (OF_PJ, '项目'), (OF_TYPE, '类型'), (OF_WHO, '对象\n（工人 / 单位）'),
                        (OF_AMT, '金额'), (OF_NOTE, '说明'), (OF_YM, '年月'), (OF_WTYPE, '对象类型'), (OF_CHK, '校验')], C_AR)
    for r in range(OF_R0, OF_R1 + 1):
        g = lambda c: f'{c}{r}'
        blank = f'AND({g(OF_DATE)}="",{g(OF_PJ)}="",{g(OF_AMT)}="")'
        ws[g(OF_SEQ)] = f'=IF({blank},"",ROW()-{OF_HDR})'
        ws[g(OF_YM)] = _ym(g(OF_DATE))
        ws[g(OF_DN)] = f'=IF(ISNUMBER({g(OF_DATE)}),INT({g(OF_DATE)}),0)'
        ws[g(OF_DN)].font = F_HELP
        ws[g(OF_WTYPE)] = f'=IF({g(OF_WHO)}="","",IFERROR(INDEX({UN_TYPES_R},MATCH({g(OF_WHO)},{UN_NAMES},0))&"",""))'
        wt = g(OF_WTYPE)
        ws[g(OF_CHK)] = (f'=IF({blank},"",IF(NOT(ISNUMBER({g(OF_DATE)})),"✗ 日期要填成日期",'
                         f'IF({g(OF_DATE)}<{OPEN_DATE},"✗ 早于建账日期",IF({g(OF_PJ)}="","✗ 选项目",IF(COUNTIF({PJ_NAMES},{g(OF_PJ)})=0,"✗ 项目不在【项目档案】里",'
                         f'IF({g(OF_TYPE)}="","✗ 选类型",IF(N({g(OF_AMT)})=0,"✗ 没有金额",'
                         f'IF(AND({g(OF_TYPE)}="总包代发工资",{wt}<>"工人",{wt}<>"管理人员"),"✗ 代发工资的对象要是一个工人（一人一行）",'
                         f'IF(AND({g(OF_TYPE)}="总包代付材料分包款",{wt}<>"材料供应商",{wt}<>"分包",{wt}<>"机械运输",{wt}<>"其他"),'
                         f'"✗ 代付款的对象要是材料商、分包或机械单位",'
                         f'"√")))))))))')
    style_rows(ws, OF_R0, OF_R1, [OF_SEQ, OF_DATE, OF_PJ, OF_TYPE, OF_WHO, OF_AMT, OF_NOTE, OF_YM, OF_WTYPE, OF_CHK],
               auto=[OF_SEQ, OF_YM, OF_WTYPE, OF_CHK], fmts={OF_DATE: DATE, OF_AMT: MONEY},
               aligns={OF_NOTE: AL, OF_CHK: AL, OF_AMT: AR},
               fills={c: FILL_IN for c in (OF_DATE, OF_PJ, OF_TYPE, OF_WHO, OF_AMT, OF_NOTE)})
    for r in range(OF_R0, OF_R1 + 1):
        ws[f'{OF_YM}{r}'].font = ws[f'{OF_WTYPE}{r}'].font = F_HELP
    _chk_cf(ws, OF_CHK, OF_R0, OF_R1)
    dv_date(ws, f'{OF_DATE}{OF_R0}:{OF_DATE}{OF_R1}')
    dv_list(ws, f'{OF_PJ}{OF_R0}:{OF_PJ}{OF_R1}', f'={PJ_NAMES}')
    dv_list(ws, f'{OF_TYPE}{OF_R0}:{OF_TYPE}{OF_R1}', '"' + ','.join(OF_TYPES) + '"')
    dv_list(ws, f'{OF_WHO}{OF_R0}:{OF_WHO}{OF_R1}', f'={UN_NAMES}', stop=False)
    hide(ws, OF_DN)
    ws.freeze_panes = f'D{OF_R0}'
    for i, row in enumerate(ctx.get('off_rows', [])):
        r = OF_R0 + i
        for c, v in zip((OF_DATE, OF_PJ, OF_TYPE, OF_WHO, OF_AMT, OF_NOTE), row):
            if v not in (None, ''):
                ws[f'{c}{r}'] = v
    return ws


def build_inv(wb, ctx):
    ws = wb.create_sheet(SH_INV)
    widths(ws, {'A': 6, 'B': 13, 'C': 11, 'D': 22, 'E': 20, 'F': 26, 'G': 20, 'H': 26, 'I': 18, 'J': 12, 'K': 11, 'L': 13,
                'M': 10, 'N': 16, 'O': 8, 'P': 8, 'Q': 8, 'R': 8, 'S': 20, 'T': 14, 'U': 7, 'V': 16, 'W': 34})
    title(ws, '发 票 登 记（开出去的销项 · 收到的进项 · A～S 跟电子税务局导出的一模一样）', IV_CHK, C_INV,
          '💡 电子税务局「发票业务 → 发票查询统计 → 全量发票查询」，开票和收票各查一次、点「导出 → 发票基础信息」，'
          '打开导出的表，从「序号」那一格起连表头下面的数据整块复制，粘到本表 A5（往下接着粘也行），列就对上了（A～S）。'
          'T 列选项目（建筑服务发票备注里有项目名；收到的材料、分包、机械的票也要选项目，不然项目税负、欠票算不准）。'
          '销方是自家公司（税号或名称对上【基础资料】①）＝销项，购方是自家公司＝进项；作废的不算；票种带「专用」的才能抵扣进项税'
          '（小规模纳税人、简易计税 3% 的项目不抵扣）。手工录也照这几列填，至少要有开票日期、销方、购方、价税合计。')
    heads = [(c, t) for c, t in zip([CL(i) for i in range(1, 20)], IV_RAW)]
    heads += [(IV_PJ, '项目\n（选）'), (IV_DIR, '方向'), (IV_UNIT, '对方单位'), (IV_CHK, '校验'), (IV_DATE, '开票日'),
              (IV_YM, '年月'), (IV_SPEC, '可抵扣'), (IV_TAXU, '税额用'), (IV_TOTU, '合计用'), (IV_RATE, '税率')]
    header(ws, IV_HDR, heads, C_INV, height=36)
    for c in (IV_PJ,):
        ws[f'{c}{IV_HDR}'].fill = fill('FFBF8F00')
    hid = (IV_DATE, IV_YM, IV_SPEC, IV_TAXU, IV_TOTU, IV_RATE)
    for c in hid:
        ws[f'{c}{IV_HDR}'].font = F_HELP
        ws[f'{c}{IV_HDR}'].fill = FILL_NONE
    me_full, me_short = f"{SH_BASE}!${CO_FULL[0]}${CO_FULL[1:]}", f"{SH_BASE}!${CO_NAME[0]}${CO_NAME[1:]}"
    me_tax = f"{SH_BASE}!${CO_TAX[0]}${CO_TAX[1:]}"
    me_type = f"{SH_BASE}!${CO_TYPE[0]}${CO_TYPE[1:]}"
    UF = rng(SH_UNIT, UN_FULL, UN_R0, UN_R1)
    UT = rng(SH_UNIT, UN_ID, UN_R0, UN_R1)
    for r in range(IV_R0, IV_R1 + 1):
        g = lambda c: f'{c}{r}'
        blank = f'AND({g(IV_SELLER)}="",{g(IV_BUYER)}="",{g(IV_TOTAL)}="",{g(IV_NET)}="")'
        isme = lambda nm, tx: f'OR(AND({tx}<>"",TRIM({tx}&"")=TRIM({me_tax}&"")),{norm(nm)}={norm(me_full)},TRIM({nm})={me_short})'
        ws[g(IV_DIR)] = f'=IF({blank},"",IF({isme(g(IV_SELLER), g(IV_STAX))},"销项",IF({isme(g(IV_BUYER), g(IV_BTAX))},"进项","？")))'
        other = f'IF({g(IV_DIR)}="销项",{g(IV_BUYER)},{g(IV_SELLER)})'
        otax = f'IF({g(IV_DIR)}="销项",{g(IV_BTAX)},{g(IV_STAX)})'
        ws[g(IV_UNIT)] = (f'=IF(OR({blank},{g(IV_DIR)}="？"),"",IFERROR(INDEX({UN_NAMES},MATCH({other},{UF},0)),'
                          f'IFERROR(INDEX({UN_NAMES},MATCH({otax}&"",{UT},0)),IFERROR(INDEX({UN_NAMES},MATCH({other},{UN_NAMES},0)),{other}&""))))')
        ws[g(IV_DATE)] = f'={date_parse(g(IV_TIME))}'
        ws[g(IV_YM)] = f'=IF(ISNUMBER({g(IV_DATE)}),YEAR({g(IV_DATE)})*100+MONTH({g(IV_DATE)}),"")'
        void = f'ISNUMBER(SEARCH("作废",{g(IV_STAT)}&""))'
        ws[g(IV_TOTU)] = (f'=IF(OR({blank},{void}),0,IF({num(g(IV_TOTAL))}<>0,{num(g(IV_TOTAL))},{num(g(IV_NET))}+{num(g(IV_TAX))}))')
        ws[g(IV_TAXU)] = f'=IF(OR({blank},{void}),0,{num(g(IV_TAX))})'
        ws[g(IV_RATE)] = f'=IF({num(g(IV_NET))}=0,"",ROUND({num(g(IV_TAX))}/{num(g(IV_NET))},2))'
        prate = f'IFERROR(INDEX({rng(SH_PROJ, PJ_TAX, PJ_R0, PJ_R1)},MATCH({g(IV_PJ)},{PJ_NAMES},0)),"")'
        simple = f'AND(ISNUMBER({prate}),N({prate})<0.05)'
        ws[g(IV_SPEC)] = (f'=IF(AND({g(IV_DIR)}="进项",ISNUMBER(SEARCH("专用",{g(IV_KIND)}&"")),{me_type}<>"小规模纳税人",NOT({simple})),1,0)')
        eno = g(IV_ENO)
        dup = (f'IF({eno}<>"",COUNTIF({ivr(IV_ENO)},"*"&{eno})>1,'
               f'AND({g(IV_NO)}<>"",COUNTIFS({ivr(IV_CODE)},{g(IV_CODE)}&"",{ivr(IV_NO)},"*"&{g(IV_NO)})>1))')
        ws[g(IV_CHK)] = (f'=IF({blank},"",IF(NOT(ISNUMBER({g(IV_DATE)})),"✗ 开票日期认不出",'
                         f'IF({g(IV_DATE)}<{OPEN_DATE},"⚠ 早于建账日期，不参与统计（建账前的开票、收票在【期初余额】填累计数）",'
                         f'IF({g(IV_DIR)}="？","✗ 销方、购方都不是自家公司（【基础资料】① 税号或全称对不上）",'
                         f'IF({void},"作废（不算）",'
                         f'IF({g(IV_TOTU)}=0,"✗ 没有金额",'
                         f'IF({g(IV_PJ)}="","⚠ 没选项目（按项目统计不到）",'
                         f'IF(COUNTIF({PJ_NAMES},{g(IV_PJ)})=0,"✗ 项目不在【项目档案】里",'
                         f'IF({dup},"⚠ 发票号码重复了（粘重了？）",'
                         f'IF(AND({g(IV_DIR)}="进项",COUNTIF({UN_NAMES},{g(IV_UNIT)})=0),"⚠ 开票单位没在【往来单位】登记（全称或税号填在往来单位就能认出）",'
                         f'"√"))))))))))')
    cols = [CL(i) for i in range(1, 20)] + [IV_PJ, IV_DIR, IV_UNIT, IV_CHK]
    style_rows(ws, IV_R0, IV_R1, cols, auto=[IV_DIR, IV_UNIT, IV_CHK],
               fmts={IV_NET: MONEY, IV_TAX: MONEY, IV_TOTAL: MONEY, IV_CODE: '@', IV_NO: '@', IV_ENO: '@', IV_STAX: '@', IV_BTAX: '@'},
               aligns={IV_SELLER: AL, IV_BUYER: AL, IV_NOTE: AL, IV_CHK: AL, IV_UNIT: AL, IV_NET: AR, IV_TAX: AR, IV_TOTAL: AR},
               fills={**{c: FILL_PASTE for c in [CL(i) for i in range(1, 20)]}, IV_PJ: FILL_IN})
    for r in range(IV_R0, IV_R1 + 1):
        for c in hid:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *hid)
    _chk_cf(ws, IV_CHK, IV_R0, IV_R1)
    dv_list(ws, f'{IV_PJ}{IV_R0}:{IV_PJ}{IV_R1}', f'={PJ_NAMES}')
    ws.auto_filter.ref = f'A{IV_HDR}:{IV_CHK}{IV_R1}'
    ws.freeze_panes = f'D{IV_R0}'
    keys = {'no': IV_NO, 'eno': IV_ENO, 'code': IV_CODE, 'stax': IV_STAX, 'seller': IV_SELLER, 'btax': IV_BTAX, 'buyer': IV_BUYER,
            'time': IV_TIME, 'net': IV_NET, 'tax': IV_TAX, 'total': IV_TOTAL, 'src': IV_SRC, 'kind': IV_KIND, 'stat': IV_STAT,
            'note': IV_NOTE, 'pj': IV_PJ}
    for i, row in enumerate(ctx['inv_rows']):
        r = IV_R0 + i
        ws[f'{IV_SEQ}{r}'] = i + 1
        for k, c in keys.items():
            v = row.get(k)
            if v not in (None, ''):
                ws[f'{c}{r}'] = v
    return ws
