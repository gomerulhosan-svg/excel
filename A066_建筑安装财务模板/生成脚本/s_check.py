# -*- coding: utf-8 -*-
"""【数据校验】：录入表逐行问题（✗ 要改、⚠ 要看）、基本信息重名/漏填、对数检查（资产负债表平不平、专户有没有余额、项目利润合计对不对得上公司利润……），
   最后列出有问题的行（表名＋行号＋问题）。首页的「提醒」引用这里的 CHK_X / CHK_W 和 CHECK_ROWS。"""
from layout import *
from common import *
from openpyxl.formatting.rule import FormulaRule

CHK_X, CHK_W = 'C4', 'F4'           # 要改（✗）合计、要看（⚠）合计
CHECK_ROWS = {}                     # 检查项 → 行号（C 列＝数，D 列＝状态）
# 其他表的核对格（各模块写好后在这里登记：(说明, 公式)；公式结果＝差额，应为 0）
EXT_DIFFS = [
    ('资产负债表 年初：资产−负债−权益', f'=ROUND(N({SH_BS}!$B$23),2)', '不为 0＝有一类钱没进资产负债表（看【资产负债表】最下面的说明）'),
    ('资产负债表 期末：资产−负债−权益', f'=ROUND(N({SH_BS}!$C$23),2)', '同上'),
    ('利润表 本期：成本费用各行之和跟总数的差', f'=ROUND(N({SH_PL}!$B$27),2)', '不为 0＝有成本费用算重了或漏了'),
    ('利润表 建账以来：成本费用各行之和跟总数的差', f'=ROUND(N({SH_PL}!$P$27),2)', '同上'),
]


def _flag(ws, r, cond_bad, kind='✗'):
    """D 列状态：cond_bad 为真＝问题"""
    put(ws, f'D{r}', f'=IF({cond_bad},"{kind}","✓")', F_TXTB, align=AC)


def build(wb, ctx=None):
    ws = wb[SH_CHK]
    widths(ws, {'A': 5, 'B': 46, 'C': 14, 'D': 7, 'E': 12, 'F': 14, 'G': 60})
    title(ws, '数 据 校 验', 'G', C_RPT,
          '💡 录完数先看这张：✗＝一定要改（不改报表会少算或算错），⚠＝看一下对不对。下面先是每张录入表有几条问题，再是对数检查，'
          '最后列出有问题的那几行（表名＋行号），回到那张表按行号找到改掉就行。这张表全是自动的，不用填。')
    home_link(ws, 'G3')
    put(ws, 'B4', '要改的（✗）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'E4', '要看的（⚠）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.row_dimensions[4].height = 24

    # ① 录入表逐行
    r = 6
    section(ws, r, 'A', 'G', '① 录入表逐行检查（每行最右边的问题汇总）', C_RPT)
    header(ws, r + 1, [('A', '#'), ('B', '表'), ('C', '✗ 要改'), ('D', ''), ('E', '⚠ 要看'), ('F', ''), ('G', '常见原因')], C_RPT)
    srcs = [(SH_CASH, '收_校验', '账户/收支项目/项目/人员 不在基本信息里，没填日期或金额，内部转账没填对方账户'),
            (SH_AR_IN, '应_校验', '项目不在项目档案、类型没选、没填金额'),
            (SH_AP_IN, '付_校验', '供应商不在供应商信息、项目不对、没填日期/金额'),
            (SH_ATT, '考_校验', '人员不在人员信息、没有工资单价、项目名不对、一个月超过 31 天')]
    xs, wsum = [], []
    for i, (sh, nm, why) in enumerate(srcs):
        rr = r + 2 + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        c = put(ws, f'B{rr}', sh, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), align=AL)
        link(c, sh)
        put(ws, f'C{rr}', f'=COUNTIF({nm},"✗*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'E{rr}', f'=COUNTIF({nm},"⚠*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'D{rr}', f'=IF(C{rr}>0,"✗","✓")', F_TXTB, align=AC)
        put(ws, f'F{rr}', f'=IF(E{rr}>0,"⚠","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', why, F_NOTE, align=ALW)
        xs.append(f'C{rr}')
        wsum.append(f'E{rr}')
        CHECK_ROWS[f'{sh}逐行'] = rr
    r = r + 2 + len(srcs) + 1

    # ② 基本信息（隐藏列 Z～AF 逐行算重名、漏填）
    section(ws, r, 'A', 'G', '② 基本信息检查', C_RPT)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '个数'), ('D', '状态'), ('E', ''), ('F', ''), ('G', '怎么改')], C_RPT)
    hid = {}

    def helper(col, n, f):
        for i in range(n):
            ws[f'{col}{i + 1}'] = f(i + 1)
            ws[f'{col}{i + 1}'].font = F_HELP
        hide(ws, col)
        return f'${col}$1:${col}${n}'

    hid['项目重名'] = helper('Z', N_PJ, lambda k: f'=IF(INDEX(项目_名称,{k})="",0,IF(COUNTIF(项目_名称,{esc(f"INDEX(项目_名称,{k})")})>1,1,0))')
    hid['供应商重名'] = helper('AA', N_SUP, lambda k: f'=IF(INDEX(供应商_名称,{k})="",0,IF(COUNTIF(供应商_名称,{esc(f"INDEX(供应商_名称,{k})")})>1,1,0))')
    hid['人员重名'] = helper('AB', N_PER, lambda k: f'=IF(INDEX(人员_姓名,{k})="",0,IF(COUNTIF(人员_姓名,{esc(f"INDEX(人员_姓名,{k})")})>1,1,0))')
    hid['账户问题'] = helper('AC', N_ACC, lambda k: (f'=IF(INDEX(账户_名称,{k})="",0,IF(OR(COUNTIF(账户_名称,{esc(f"INDEX(账户_名称,{k})")})>1,'
                                                     f'AND(INDEX(账户_类型,{k})="个人户",INDEX(账户_所属人,{k})="")),1,0))'))
    hid['收入项目'] = helper('AD', N_INC, lambda k: (f'=IF(INDEX(收入项目_名称,{k})="",0,IF(OR(INDEX(收入项目_归类,{k})="",'
                                                    f'ISNA(MATCH(INDEX(收入项目_归类,{k}),{{"' + '","'.join(CAT_NAMES) + f'"}},0))),1,0))'))
    hid['支出项目'] = helper('AE', N_EXP, lambda k: (f'=IF(INDEX(支出项目_名称,{k})="",0,IF(OR(INDEX(支出项目_归类,{k})="",'
                                                    f'ISNA(MATCH(INDEX(支出项目_归类,{k}),{{"' + '","'.join(CAT_NAMES) + f'"}},0))),1,0))'))
    hid['没开工日'] = helper('AF', N_PJ, lambda k: f'=IF(AND(INDEX(项目_名称,{k})<>"",INDEX(项目_类型,{k})="工程",INDEX(项目_开工,{k})=0),1,0)')
    base_checks = [
        ('项目档案：项目名称重复', f'=SUM({hid["项目重名"]})', '✗', '同一个项目只留一行（下拉和报表按名称认，重复了数会算到第一个）'),
        ('供应商信息：名称重复', f'=SUM({hid["供应商重名"]})', '✗', '同一家只留一行；别名写在「开票单位」或备注'),
        ('人员信息：姓名重复', f'=SUM({hid["人员重名"]})', '✗', '同名的两个人在名字后面加区分，比如「张三（电工）」'),
        ('基本信息②账户：名称重复或个人户没填所属人', f'=SUM({hid["账户问题"]})', '✗', '个人户一定要填是谁的（个人往来按人汇总）'),
        ('基本信息③收入项目：没选归类或归类不认识', f'=SUM({hid["收入项目"]})', '✗', '在 H 列下拉选一个归类'),
        ('基本信息④支出项目：没选归类或归类不认识', f'=SUM({hid["支出项目"]})', '✗', '在 K 列下拉选一个归类'),
        ('项目档案：工程项目没填开工日期', f'=SUM({hid["没开工日"]})', '⚠', '没填按建账日算「开工至今」；填上开工日期项目账更准'),
    ]
    r += 2
    for i, (lab, f, kind, how) in enumerate(base_checks):
        rr = r + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        put(ws, f'B{rr}', lab, F_TXT, align=AL)
        put(ws, f'C{rr}', f, F_AUTOB, fmt='0', align=AC)
        _flag(ws, rr, f'C{rr}>0', kind)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        (xs if kind == '✗' else wsum).append(f'IF(C{rr}>0,1,0)')
        CHECK_ROWS[lab] = rr
    r += len(base_checks) + 1

    # ③ 对数检查
    section(ws, r, 'A', 'G', '③ 对数检查（金额；差额应为 0）', C_RPT)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '金额/个数'), ('D', '状态'), ('E', ''), ('F', ''), ('G', '说明')], C_RPT)
    acc = lambda c: f'{H_SUM}!${c}${SUM_ACC0}:${c}${SUM_ACC0 + N_ACC - 1}'
    sup = f'{H_SUM}!$D${SUM_SUP0}:$D${SUM_SUP0 + N_SUP - 1}'
    num_checks = [
        ('专户（农民工专户、九安账户这类）期末余额合计', f'=ROUND(SUMIFS({acc("E")},{acc("B")},"专户"),2)', 'ABS(C{r})>=0.01', '⚠',
         '专户应该一进一出余额为 0；不为 0＝有一边漏记（收了没付出去，或付了没记收）。到【资金报表】看是哪个专户'),
        ('专户里余额不为 0 的个数', f'=COUNTIFS({acc("B")},"专户",{acc("E")},">=0.01")+COUNTIFS({acc("B")},"专户",{acc("E")},"<=-0.01")',
         'C{r}>0', '⚠', '同上'),
        ('内部转账没填对方账户（或对方账户不对）的笔数', '=COUNTIFS(收_归类原,"内部转账",收_归类,"",收_有效,1)', 'C{r}>0', '✗',
         '这几笔两边账户对不上，按「未分类」进了利润表；在【收支登记】N 列选对方账户'),
        ('没认出归类的收支（未分类）笔数', '=COUNTIFS(收_成本类,"未分类")', 'C{r}>0', '✗', '收支项目不在基本信息③④里，或者没选归类；金额见下一行'),
        ('没认出归类的收支（未分类）净额', '=ROUND(SUMIFS(收_净额,收_成本类,"未分类"),2)', 'ABS(C{r})>=0.01', '✗', '利润表「未分类收支」就是这个数'),
        ('待摊费用没摊出去的（各年合计）', '=ROUND(SUM(汇_未摊),2)', 'ABS(C{r})>=0.01', '⚠',
         '那一年没有项目有施工费（分摊基数为 0），待摊费用留在公司；看【费用分摊】'),
        ('供应商多付了（应付余额为负）的家数', f'=COUNTIF({sup},"<=-0.01")', 'C{r}>0', '⚠',
         '付款比应付登记多：可能是送货单/结算没录全，或者付款的供应商名称填错了；看【应付账款】未付为负的那几家'),
        ('材料/分包/机械付款没对上应付登记（直接进成本）的笔数', '=COUNTIFS(收_成本类,"材料")+COUNTIFS(收_成本类,"分包")+COUNTIFS(收_成本类,"机械")',
         'C{r}>0', '⚠', '供应商名称空着，或者这家在【应付登记】里一行都没有；零星采购可以不管，大额的建议补应付登记'),
        ('工资发给了考勤里没有的人（直接算人工成本）的笔数',
         '=COUNTIFS(收_归类,"工资发放",收_冲工资,0,收_过账,0,收_日期,">="&P_考勤起算)', 'C{r}>0', '⚠',
         '人员没填或者这个人没录考勤：这笔工资直接算成本，不冲欠薪'),
        ('期初余额表里填了「未分配利润」', '=COUNTIF(期初_类型,"未分配利润")', 'C{r}>0', '⚠',
         '期初未分配利润由系统倒挤（期初资产−负债−实收资本），填的这个数不用'),
    ] + [(lab, f, 'ABS(C{r})>=0.01', '✗', how) for lab, f, how in EXT_DIFFS]
    r += 2
    for i, (lab, f, bad, kind, how) in enumerate(num_checks):
        rr = r + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        put(ws, f'B{rr}', lab, F_TXT, align=ALW)
        put(ws, f'C{rr}', f, F_AUTOB, fmt=MONEY if '额' in lab or '差' in lab or '合计' in lab else '0', align=AR)
        _flag(ws, rr, bad.format(r=rr), kind)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 30
        (xs if kind == '✗' else wsum).append(f'IF({bad.format(r=rr)},1,0)')
        CHECK_ROWS[lab] = rr
    r += len(num_checks) + 1
    put(ws, CHK_X, '=' + '+'.join(xs), F_KPI_V, fill('FFFFFFFF'), '0', AC)
    put(ws, CHK_W, '=' + '+'.join(wsum), F_KPI_V, fill('FFFFFFFF'), '0', AC)

    # ④ 有问题的行（每张录入表最多 150 行）
    section(ws, r, 'A', 'G', '④ 有问题的行（回到那张表按行号找）', C_RPT)
    header(ws, r + 1, [('A', '#'), ('B', '问题'), ('C', '表'), ('D', '行号'), ('E', '日期'), ('F', '金额'), ('G', '摘要 / 名称')], C_RPT)
    r += 2
    lists = [(SH_CASH, '收', N_CASH, '收_校验', '收_录入行', '收_日期', 'INDEX(收_净额原,{k})', 'INDEX(收_收支项目,{k})&"｜"&INDEX(收_摘要,{k})'),
             (SH_AR_IN, '应', N_AR, '应_校验', '应_录入行', '应_日期', 'INDEX(应_金额,{k})', 'INDEX(应_项目,{k})&"｜"&INDEX(应_类型,{k})'),
             (SH_AP_IN, '付', N_AP, '付_校验', '付_录入行', '付_日期', 'INDEX(付_应付额,{k})', 'INDEX(付_供应商,{k})&"｜"&INDEX(付_摘要,{k})'),
             (SH_ATT, '考', N_ATT, '考_校验', '考_录入行', '考_月', 'INDEX(考_应发,{k})', 'INDEX(考_姓名,{k})')]
    SHOW = 150
    cnt_cols = iter(['AH', 'AI', 'AJ', 'AK'])
    for sh, tag, n, chk, row_nm, date_nm, amt, desc in lists:
        cnt_c = next(cnt_cols)
        counter(ws, cnt_c, 1, n, lambda i, chk=chk: f'INDEX({chk},{i + 1})<>""')
        hide(ws, cnt_c)
        total = f'${cnt_c}${n}'
        put(ws, f'B{r}', f'="{sh}：共 "&{total}&" 行有问题"&IF({total}>{SHOW},"（只列前 {SHOW} 行）","")', F_TXTB, FILL_SUB, align=AL)
        ws.merge_cells(f'B{r}:G{r}')
        r += 1
        for k in range(1, SHOW + 1):
            rr = r + k - 1
            ix = f'$A{rr}'
            ws[f'A{rr}'] = f'=IF({kth(k, cnt_c, 1, n)}=0,"",{kth(k, cnt_c, 1, n)})'
            ws[f'A{rr}'].font = F_HELP
            g = lambda e: f'=IF({ix}="","",{e.format(k=ix)})'
            put(ws, f'B{rr}', g(f'INDEX({chk},{{k}})'), F_TXT, align=AL)
            put(ws, f'C{rr}', f'=IF({ix}="","","{sh}")', F_TXT, align=AC)
            put(ws, f'D{rr}', g(f'INDEX({row_nm},{{k}})'), F_TXTB, fmt='0', align=AC)
            put(ws, f'E{rr}', g(f'IF(INDEX({date_nm},{{k}})=0,"",INDEX({date_nm},{{k}}))'), F_TXT, fmt=DATE, align=AC)
            put(ws, f'F{rr}', g(amt), F_TXT, fmt=MONEY, align=AR)
            put(ws, f'G{rr}', g(desc), F_NOTE, align=AL)
        r += SHOW + 1
    ws.conditional_formatting.add(f'D1:D{r}', FormulaRule(formula=['D1="✗"'], font=F_RED))
    ws.conditional_formatting.add(f'F1:F12', FormulaRule(formula=['F1="⚠"'], font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.conditional_formatting.add(f'D1:D{r}', FormulaRule(formula=['D1="⚠"'], font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:4', landscape=False)
    ws.sheet_view.showGridLines = False
