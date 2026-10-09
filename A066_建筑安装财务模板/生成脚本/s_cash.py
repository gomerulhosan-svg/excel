# -*- coding: utf-8 -*-
"""查看表（绿）：
   【资金报表】出纳报表：周报 / 月报 / 年报 / 自定义。① 各账户期初、收入、支出、期末（按类型小计）；② 按收支项目（本期、本年累计）；
              ③ 按项目（收工程款、其他收入、支出、净额）；④ 年报：各账户每月月末余额、收支项目 × 12 个月。
   【账户明细】选账户（空＝全部）＋起止：按日期逐笔列出，余额逐笔滚动。
   口径（报表口径.md §2）：账户余额＝期初＋本账户净额−（内部转账里「对方账户」是它的那些行的净额）；
   内部转账只记一行：账户＝转出/转入，对方账户＝另一个；对方那边收入、支出对调。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from layout import *
from common import *

GREEN_H = 'FF70AD47'                      # 表头浅一点的绿
F_GRAY_I = Font(name=YH, sz=10, italic=True, color='FF808080')
F_RED_CF = Font(color='FFC00000', bold=True)
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'    # 零不显示（明细清单用）
CAT_ARR = '{"' + '","'.join(CAT_NAMES) + '"}'


def per(a, b):
    return f'收_日期,">="&{a},收_日期,"<="&{b}'


def acc_inc(e, a, b):
    """账户（名字条件 e）在 [a,b] 的收入：本账户收入栏 ＋ 内部转账里对方账户是它的那些行的支出栏"""
    return f'SUMIFS(收_收入,收_账户,{e},{per(a, b)})+SUMIFS(收_支出,收_对方账户,{e},收_归类,"内部转账",{per(a, b)})'


def acc_exp(e, a, b):
    return f'SUMIFS(收_支出,收_账户,{e},{per(a, b)})+SUMIFS(收_收入,收_对方账户,{e},收_归类,"内部转账",{per(a, b)})'


def acc_bal(i, e, crit):
    """第 i 个账户（e＝名字条件）在 收_日期 满足 crit（如 "<="&d）时的余额"""
    return (f'INDEX(账户_期初,{i})+SUMIFS(收_净额,收_账户,{e},收_日期,{crit})'
            f'-SUMIFS(收_净额,收_对方账户,{e},收_归类,"内部转账",收_日期,{crit})')


def cf_warn(ws, rng, first):
    """「⚠」开头的说明标红"""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], font=F_RED_CF))


def cell(ws, coord, v, fmt=None, align=AR, font=F_TXT, fill_=None):
    return put(ws, coord, v, font, fill_, fmt, align)


# ═════════════════════════════════════════ 资金报表 ═════════════════════════════════════════
LAST = 'Q'
HR0 = 5                                                     # 隐藏计算块的第一行（第 i 个账户/项目在 HR0+i-1 行）
H = dict(名称='AB', 类型='AC', 所属人='AD', 组='AE', 期初='AF', 收入='AG', 支出='AH', 期末='AI', 笔数='AJ', 年初='AK',
         全年收入='AX', 全年支出='AY')
MCOL = [CL(CI('AL') + m) for m in range(12)]                # AL..AW：1～12 月末余额（第 1 行月初、第 2 行有效月末、第 3 行显示否）
VMON = [CL(CI('D') + m) for m in range(12)]                 # 显示：D..O
GCNT = {1: 'AZ', 2: 'BA', 3: 'BB'}                          # 账户分组计数：1 银行＋现金（及其他），2 专户，3 个人户
ICNT, ECNT = 'BC', 'BD'                                     # 收入项目、支出项目计数
PACT, PCA, PCB = 'BE', 'BF', 'BG'                           # 项目：本期笔数（-1＝空名）、有收支的计数、没收支的计数
KEYLBL = 'BH'
IDX = 'Z'                                                   # 显示行用的「第几个」
SC = 'AA'                                                   # 标量
S0, S1 = '$B$5', '$D$5'                                     # 本期起、止
Y0, Y1 = f'${SC}$4', f'${SC}$5'                             # 年报的年初、年末（不超过截止日）
YB0 = f'${SC}$20'                                           # ② 本年累计的起：本期止那一年的 1 月 1 日
TYP, DEF, YR, XFER, YPREV, NEND, YXFER = (f'${SC}${i}' for i in (1, 2, 3, 6, 7, 8, 9))
GROUPS_ = {1: '银行＋现金', 2: '专户', 3: '个人户'}


def hrng(col, n=N_ACC):
    return f'${col}${HR0}:${col}${HR0 + n - 1}'


def _grp(t):
    return 2 if t == '专户' else (3 if t == '个人户' else 1)


def build_fund(ws, ctx):
    accs = [a for a in ctx.get('accounts', []) if str(a.get('名称', '')).strip()]
    ng = {g: sum(1 for a in accs if _grp(str(a.get('类型', '')).strip()) == g) for g in (1, 2, 3)}
    M = {1: min(N_ACC, ng[1] + 1), 2: min(N_ACC, ng[2] + 1), 3: min(N_ACC, ng[3] + 2)}
    M_INC = min(N_INC, len(ctx.get('inc_items', [])) + 3)
    M_EXP = min(N_EXP, len(ctx.get('exp_items', [])) + 5)
    M_PJ = min(N_PJ, len(ctx.get('projects', [])) + 3)

    widths(ws, {'A': 20, 'B': 11, **{CL(i): 13 for i in range(3, 18)}})
    tip = ('💡 出纳报表。黄格选「报表类型」：周报＝日期所在那一周（周一～周日）、月报＝那个月、年报＝那一年、自定义＝自己填起止；'
           '「日期」空着＝首页截止日。下面自动出：① 每个账户的期初、收入、支出、期末余额；② 按收支项目；③ 按项目；④ 这一年每月月末余额和每月收支（给老板看的年报）。'
           '账户之间倒钱（内部转账）在①里转出、转入两边各算一次，②③不算收支。专户（农民工专户等）余额应该是 0；'
           '个人户是老板、员工的个人账户：负数＝他替公司垫的钱，正数＝他手上有公司的钱，不算货币资金。')
    title(ws, '资金报表（出纳周报 / 月报 / 年报）', LAST, C_VIEW, tip)

    # ── 选择格 ──
    selector(ws, 'A3', '报表类型', 'B3', '月报', '"周报,月报,年报,自定义"',
             prompt='周报＝日期所在那一周；月报＝那个月；年报＝那一年；自定义＝填右边的起止')
    selector(ws, 'C3', '日期', 'D3', None, fmt=DATE, prompt='空＝首页截止日')
    selector(ws, 'E3', '自定义起', 'F3', None, fmt=DATE, prompt='报表类型选「自定义」才用；空＝首页年初')
    selector(ws, 'G3', '自定义止', 'H3', None, fmt=DATE, prompt='报表类型选「自定义」才用；空＝首页截止日')
    for c in ('D3', 'F3', 'H3'):
        dv_date(ws, c)
    home_link(ws, f'{LAST}3')
    put(ws, 'A4', '实际用的', F_NOTE, align=AC)
    cell(ws, 'B4', f'={TYP}', align=AC, font=F_AUTOB)
    put(ws, 'C4', '空＝首页截止日→', F_NOTE, align=AR)
    cell(ws, 'D4', f'={DEF}', DATE, AC, F_AUTOB)
    put(ws, 'E4', '空＝首页年初→', F_NOTE, align=AR)
    cell(ws, 'F4', '=IF(ISNUMBER(F3),INT(F3),P_年初)', DATE, AC, F_AUTOB)
    put(ws, 'G4', '空＝首页截止日→', F_NOTE, align=AR)
    cell(ws, 'H4', '=IF(ISNUMBER(H3),INT(H3),P_截止)', DATE, AC, F_AUTOB)
    ws.merge_cells(f'I4:{LAST}4')
    put(ws, 'I4', '自定义起止只在报表类型选「自定义」时用；周报、月报、年报都到首页截止日为止', F_NOTE, align=AL, border=False)
    put(ws, 'A5', '本期', F_KPI_L, fill('FFD9E1F2'), align=AC)
    cell(ws, 'B5', (f'=IF({TYP}="周报",{DEF}-WEEKDAY({DEF},2)+1,IF({TYP}="月报",DATE(YEAR({DEF}),MONTH({DEF}),1),'
                    f'IF({TYP}="年报",DATE(YEAR({DEF}),1,1),F4)))'), DATE, AC, F_KPI_V)
    put(ws, 'C5', '至', F_KPI_L, align=AC)
    cell(ws, 'D5', f'=IF(AND({TYP}<>"自定义",P_截止>={S0},P_截止<{NEND}),P_截止,{NEND})', DATE, AC, F_KPI_V)
    ws.merge_cells(f'E5:{LAST}5')
    put(ws, 'E5', (f'=IF({TYP}="周报","周报：日期所在那一周（周一～周日）",IF({TYP}="月报","月报：日期所在的那个月",'
                   f'IF({TYP}="年报","年报：日期所在那一年（1 月 1 日起）","自定义：用上面的自定义起、止")))'
                   f'&IF({S0}>{S1},"　⚠ 起晚于止，请改日期","")'), F_AUTOB, align=AL, border=False)
    ws.row_dimensions[5].height = 22

    # ── 标量（隐藏 AA 列） ──
    sc = {
        1: '=IF(OR(TRIM(B3&"")="周报",TRIM(B3&"")="月报",TRIM(B3&"")="年报",TRIM(B3&"")="自定义"),TRIM(B3&""),"月报")',
        2: '=IF(ISNUMBER(D3),INT(D3),P_截止)',
        3: f'=IF({TYP}="自定义",YEAR({S1}),YEAR({DEF}))',
        4: f'=DATE({YR},1,1)',
        5: f'=MIN(DATE({YR},12,31),P_截止)',
        6: f'=SUMIFS(收_收入,收_归类,"内部转账",{per(S0, S1)})+SUMIFS(收_支出,收_归类,"内部转账",{per(S0, S1)})',
        7: f'={Y0}-1',
        8: (f'=IF({TYP}="周报",{S0}+6,IF({TYP}="月报",DATE(YEAR({DEF}),MONTH({DEF})+1,0),'
            f'IF({TYP}="年报",DATE(YEAR({DEF}),12,31),H4)))'),
        9: f'=SUMIFS(收_收入,收_归类,"内部转账",{per(Y0, Y1)})+SUMIFS(收_支出,收_归类,"内部转账",{per(Y0, Y1)})',
        20: f'=DATE(YEAR({S1}),1,1)',
    }
    for i, f in sc.items():
        ws[f'{SC}{i}'] = f
        ws[f'{SC}{i}'].font = F_HELP
    for m in range(12):
        c = MCOL[m]
        ws[f'{c}1'] = f'=DATE({YR},{m + 1},1)'
        ws[f'{c}2'] = f'=MIN(DATE({YR},{m + 2},0),P_截止)'
        ws[f'{c}3'] = f'=IF({c}1<=P_截止,1,0)'

    # ── 隐藏计算块：每个账户（N_ACC 个位置） ──
    for i in range(N_ACC):
        r, n = HR0 + i, i + 1
        a = f'${H["名称"]}{r}'
        e = esc(a)
        f = {
            '名称': f'=INDEX(账户_名称,{n})&""',
            '类型': f'=INDEX(账户_类型,{n})&""',
            '所属人': f'=INDEX(账户_所属人,{n})&""',
            '组': f'=IF({a}="",0,IF({H["类型"]}{r}="专户",2,IF({H["类型"]}{r}="个人户",3,1)))',
            '期初': f'=IF({a}="",0,ROUND({acc_bal(n, e, chr(34) + "<" + chr(34) + "&" + S0)},2))',
            '收入': f'=IF({a}="",0,ROUND({acc_inc(e, S0, S1)},2))',
            '支出': f'=IF({a}="",0,ROUND({acc_exp(e, S0, S1)},2))',
            '期末': f'=ROUND({H["期初"]}{r}+{H["收入"]}{r}-{H["支出"]}{r},2)',
            '笔数': (f'=IF({a}="",0,COUNTIFS(收_账户,{e},收_有效,1,{per(S0, S1)})'
                   f'+COUNTIFS(收_对方账户,{e},收_归类,"内部转账",{per(S0, S1)}))'),
            '年初': f'=IF({a}="",0,ROUND({acc_bal(n, e, chr(34) + "<=" + chr(34) + "&" + YPREV)},2))',
            '全年收入': f'=IF({a}="",0,ROUND({acc_inc(e, Y0, Y1)},2))',
            '全年支出': f'=IF({a}="",0,ROUND({acc_exp(e, Y0, Y1)},2))',
        }
        for k, v in f.items():
            ws[f'{H[k]}{r}'] = v
        for m in range(12):
            c = MCOL[m]
            ws[f'{c}{r}'] = f'=IF({a}="",0,ROUND({acc_bal(n, e, chr(34) + "<=" + chr(34) + "&" + c + "$2")},2))'
    for g, col in GCNT.items():
        counter(ws, col, HR0, N_ACC, lambda i, g=g: f'${H["组"]}{HR0 + i}={g}')
    counter(ws, ICNT, HR0, N_INC, lambda i: f'INDEX(收入项目_名称,{i + 1})<>""')
    counter(ws, ECNT, HR0, N_EXP, lambda i: f'INDEX(支出项目_名称,{i + 1})<>""')
    for i in range(N_PJ):
        r = HR0 + i
        nm = f'INDEX(项目_名称,{i + 1})'
        ws[f'{PACT}{r}'] = (f'=IF({nm}="",-1,COUNTIFS(收_项目原,{esc(nm)},收_有效,1,收_归类,"<>内部转账",{per(S0, S1)}))')
    counter(ws, PCA, HR0, N_PJ, lambda i: f'${PACT}{HR0 + i}>0')
    counter(ws, PCB, HR0, N_PJ, lambda i: f'${PACT}{HR0 + i}=0')
    gcnt = {g: cnt(c, HR0, N_ACC) for g, c in GCNT.items()}
    n_inc, n_exp = cnt(ICNT, HR0, N_INC), cnt(ECNT, HR0, N_EXP)
    n_pa, n_pb = cnt(PCA, HR0, N_PJ), cnt(PCB, HR0, N_PJ)

    def hidx(r, formula):
        ws[f'{IDX}{r}'] = '=' + formula
        ws[f'{IDX}{r}'].font = F_HELP
        return f'${IDX}{r}'

    def hv(col, z):
        return f'INDEX({hrng(col)},{z})'

    def style_line(r, cols, kind=None, fmts=None, aligns=None):
        fmts, aligns = fmts or {}, aligns or {}
        fl = {'sub': FILL_SUB, 'tot': FILL_TOT, 'head': fill('FFE2EFDA')}.get(kind)
        for c in cols:
            x = ws[f'{c}{r}']
            x.font = F_AUTOB if kind else F_TXT
            x.border = BD
            if fl:
                x.fill = fl
            x.number_format = fmts.get(c, MONEY)
            x.alignment = aligns.get(c, AR)

    # ═════ ① 各账户 ═════
    r = 7
    section(ws, r, 'A', 'I', '① 各账户本期收支（收入、支出含账户之间倒的钱；期初＝起日前一天的余额）', C_VIEW)
    r += 1
    header(ws, r, [('A', '账户'), ('B', '类型'), ('C', '所属人'), ('D', '期初余额\n（起日前一天）'), ('E', '本期收入'), ('F', '本期支出'),
                   ('G', '期末余额'), ('H', '本期笔数'), ('I', '说明')], GREEN_H)
    r += 1
    first_acc = r
    sub = {}
    acc_cols = list('ABCDEFGHI')
    afmt = {'A': 'General', 'B': 'General', 'C': 'General', 'H': INT, 'I': 'General'}
    aal = {'A': AL, 'B': AC, 'C': AC, 'H': AC, 'I': AL}
    for g in (1, 2, 3):
        for k in range(1, M[g] + 1):
            z = hidx(r, kth(k, GCNT[g], HR0, N_ACC))
            for c, col in (('A', '名称'), ('B', '类型'), ('C', '所属人')):
                ws[f'{c}{r}'] = f'=IF({z}=0,"",{hv(H[col], z)})'
            for c, col in (('D', '期初'), ('E', '收入'), ('F', '支出'), ('G', '期末'), ('H', '笔数')):
                ws[f'{c}{r}'] = f'=IF({z}=0,"",{hv(H[col], z)})'
            ws[f'I{r}'] = (f'=IF({z}=0,"",IF(B{r}="专户",IF(ROUND(G{r},2)<>0,"⚠ 专户有钱没付出去或漏记一边","专户余额为 0 ✓"),'
                           f'IF(B{r}="个人户",IF(G{r}<0,"他替公司垫的钱（公司还了多少看【个人往来】）",IF(G{r}>0,"他手上有公司的钱","")),'
                           f'IF(G{r}<0,"⚠ 余额是负数：漏记收入或记错账户？",""))))')
            style_line(r, acc_cols, None, afmt, aal)
            r += 1
        sub[g] = r
        gs = lambda col, g=g: f'SUMIFS({hrng(H[col])},{hrng(H["组"])},{g})'
        ws[f'A{r}'] = {1: '银行＋现金 小计（货币资金）', 2: '专户 小计（应为 0）', 3: '个人户 小计（不算货币资金）'}[g]
        for c, col in (('D', '期初'), ('E', '收入'), ('F', '支出'), ('G', '期末'), ('H', '笔数')):
            ws[f'{c}{r}'] = f'={gs(col)}'
        over = f'IF({gcnt[g]}>{M[g]},"　⚠ 共 "&{gcnt[g]}&" 个，上面只列了前 {M[g]} 个（小计是全部的）","")'
        if g == 1:
            ws[f'I{r}'] = f'="共 "&{gcnt[g]}&" 个账户"&{over}'
        elif g == 2:
            ws[f'I{r}'] = f'=IF({gcnt[g]}=0,"没有专户",IF(ROUND(G{r},2)<>0,"⚠ 专户有钱没付出去或漏记一边","✓ 专户余额为 0"))&{over}'
        else:
            ws[f'I{r}'] = f'=IF({gcnt[g]}=0,"没有个人户","负数＝公司欠他们（垫付的；公司已还的报销款看【个人往来】），正数＝他们手上有公司的钱")&{over}'
        style_line(r, acc_cols, 'sub', afmt, {**aal, 'A': AL})
        r += 1
    tot_acc = r
    ws[f'A{r}'] = '全部账户 合计'
    for c, col in (('D', '期初'), ('E', '收入'), ('F', '支出'), ('G', '期末'), ('H', '笔数')):
        ws[f'{c}{r}'] = f'=SUM({hrng(H[col])})'
    ws[f'I{r}'] = (f'=IF(ROUND(G{sub[1]}+G{sub[2]}+G{sub[3]}-G{r},2)<>0,"⚠ 有账户类型不是 银行/现金/专户/个人户",'
                   f'"＝ 银行＋现金 ＋ 专户 ＋ 个人户")')
    style_line(r, acc_cols, 'tot', afmt, {**aal, 'A': AL})
    r += 1
    xfer_row = r
    ws[f'A{r}'] = '其中：账户之间倒钱'
    ws[f'E{r}'] = f'={XFER}'
    ws[f'F{r}'] = f'={XFER}'
    ws[f'I{r}'] = '内部转账一转一收，两边各算一次，不是真的收支（②③都不算）'
    style_line(r, acc_cols, None, afmt, {**aal, 'A': AL})
    for c in acc_cols:
        ws[f'{c}{r}'].font = F_GRAY_I
    cf_warn(ws, f'I{first_acc}:I{r}', f'I{first_acc}')
    r += 2

    # ═════ ② 按收支项目 ═════
    section(ws, r, 'A', 'I', '② 按收支项目（本期＝上面的起止；本年累计＝本期止那一年的 1 月 1 日到本期止）', C_VIEW)
    r += 1
    header(ws, r, [('A', '收支项目'), ('B', '归类'), ('C', '本期金额'), ('D', '本年累计\n（1 月 1 日～本期止）'), ('E', '说明')], GREEN_H)
    r += 1
    it_cols = list('ABCDE')
    ifmt = {'A': 'General', 'B': 'General', 'E': 'General'}
    ial = {'A': AL, 'B': AC, 'E': AL}
    yr0, yr1 = YB0, S1
    blocks = {}
    for side, label, cntc, nlist, mshow, nm_name, cat_name, sign, crit in (
            ('收入', '收入项目', ICNT, N_INC, M_INC, '收入项目_名称', '收入项目_归类', '', '"收入"'),
            ('支出', '支出项目', ECNT, N_EXP, M_EXP, '支出项目_名称', '支出项目_归类', '-', '"<>收入"')):
        ws[f'A{r}'] = label
        style_line(r, it_cols, 'head', ifmt, ial)
        r += 1
        i0 = r
        for k in range(1, mshow + 1):
            z = hidx(r, kth(k, cntc, HR0, nlist))
            e = esc(f'$A{r}')
            ws[f'A{r}'] = f'=IF({z}=0,"",INDEX({nm_name},{z}))'
            ws[f'B{r}'] = f'=IF({z}=0,"",INDEX({cat_name},{z}))'
            ws[f'C{r}'] = f'=IF($A{r}="","",{sign}SUMIFS(收_净额,收_收支项目,{e},收_类别,{crit},{per(S0, S1)}))'
            ws[f'D{r}'] = f'=IF($A{r}="","",{sign}SUMIFS(收_净额,收_收支项目,{e},收_类别,{crit},{per(yr0, yr1)}))'
            ws[f'E{r}'] = (f'=IF($A{r}="","",IF(B{r}="内部转账","账户间倒钱，不算收支（不进合计）",'
                           f'IF(ISNA(MATCH(B{r},{CAT_ARR},0)),"⚠ 这个收支项目没选归类（【基本信息】）","")))')
            style_line(r, it_cols, None, ifmt, ial)
            r += 1
        i1 = r - 1
        tot_r = r + 1
        ws[f'A{r}'] = f'其他没对上的{side}'
        for c in 'CD':
            ws[f'{c}{r}'] = f'=ROUND({c}{tot_r}-SUMIFS({c}{i0}:{c}{i1},$B{i0}:$B{i1},"<>内部转账"),2)'
        ws[f'E{r}'] = (f'=IF(AND(ROUND(C{r},2)=0,ROUND(D{r},2)=0),"",'
                       f'"收支项目不在清单里、类别选反、或内部转账没填对方账户（看【收支登记】校验列）")'
                       f'&IF({cnt(cntc, HR0, nlist)}>{mshow},"　⚠ 共 "&{cnt(cntc, HR0, nlist)}&" 个{label}，只列了前 {mshow} 个","")')
        style_line(r, it_cols, None, ifmt, ial)
        r += 1
        ws[f'A{r}'] = f'{side}合计（不含内部转账）'
        ws[f'C{r}'] = f'={sign}SUMIFS(收_净额,收_类别,{crit},收_归类,"<>内部转账",{per(S0, S1)})'
        ws[f'D{r}'] = f'={sign}SUMIFS(收_净额,收_类别,{crit},收_归类,"<>内部转账",{per(yr0, yr1)})'
        style_line(r, it_cols, 'tot', ifmt, ial)
        blocks[side] = (i0, i1, r)
        ws.conditional_formatting.add(f'A{i0}:D{i1}', FormulaRule(formula=[f'$B{i0}="内部转账"'], font=Font(color='FF808080', italic=True)))
        cf_warn(ws, f'E{i0}:E{r}', f'E{i0}')
        r += 1
    net_row = r
    ti, te = blocks['收入'][2], blocks['支出'][2]
    ws[f'A{r}'] = '净增加（收入合计−支出合计）'
    ws[f'C{r}'] = f'=ROUND(C{ti}-C{te},2)'
    ws[f'D{r}'] = f'=ROUND(D{ti}-D{te},2)'
    ws[f'E{r}'] = (f'=IF(ROUND(C{r}-(G{tot_acc}-D{tot_acc}),2)=0,"✓ 本期净增加＝①全部账户期末−期初",'
                   f'"⚠ 跟①全部账户余额的变化差 "&ROUND(C{r}-(G{tot_acc}-D{tot_acc}),2))')
    style_line(r, it_cols, 'tot', ifmt, ial)
    cf_warn(ws, f'E{r}', f'E{r}')
    r += 2

    # ═════ ③ 按项目 ═════
    section(ws, r, 'A', 'I', '③ 按项目（本期；有收支的项目排在前面；办公室这类公司项目也列）', C_VIEW)
    r += 1
    header(ws, r, [('A', '项目'), ('B', '类型'), ('C', '收工程款'), ('D', '其他收入'), ('E', '支出合计'), ('F', '净额\n（收−支）'),
                   ('G', '本期笔数')], GREEN_H)
    r += 1
    pj_cols = list('ABCDEFG')
    pfmt = {'A': 'General', 'B': 'General', 'G': INT}
    pal = {'A': AL, 'B': AC, 'G': AC}
    p0 = r
    gk = '收_归类,"<>工程款收款",收_归类,"<>内部转账"'
    for k in range(1, M_PJ + 1):
        z = hidx(r, f'IF({k}<={n_pa},{kth(k, PCA, HR0, N_PJ)},{kth(f"({k}-{n_pa})", PCB, HR0, N_PJ)})')
        e = esc(f'$A{r}')
        ws[f'A{r}'] = f'=IF({z}=0,"",INDEX(项目_名称,{z}))'
        ws[f'B{r}'] = f'=IF({z}=0,"",INDEX(项目_类型,{z}))'
        ws[f'C{r}'] = f'=IF($A{r}="","",SUMIFS(收_净额,收_项目原,{e},收_归类,"工程款收款",{per(S0, S1)}))'
        ws[f'D{r}'] = f'=IF($A{r}="","",SUMIFS(收_净额,收_项目原,{e},收_类别,"收入",{gk},{per(S0, S1)}))'
        ws[f'E{r}'] = f'=IF($A{r}="","",-SUMIFS(收_净额,收_项目原,{e},收_类别,"<>收入",{gk},{per(S0, S1)}))'
        ws[f'F{r}'] = f'=IF($A{r}="","",ROUND(C{r}+D{r}-E{r},2))'
        ws[f'G{r}'] = f'=IF($A{r}="","",INDEX({hrng(PACT, N_PJ)},{z}))'
        style_line(r, pj_cols, None, pfmt, pal)
        r += 1
    p1 = r - 1
    ptot = r + 1
    ws[f'A{r}'] = '没填项目 / 项目不在档案里'
    for c in 'CDEFG':
        ws[f'{c}{r}'] = f'=ROUND({c}{ptot}-SUM({c}{p0}:{c}{p1}),2)'
    style_line(r, pj_cols, None, pfmt, pal)
    ws[f'H{r}'] = f'=IF({n_pa}+{n_pb}>{M_PJ},"⚠ 共 "&({n_pa}+{n_pb})&" 个项目，只列了前 {M_PJ} 个（合计是全部的）","")'
    ws[f'H{r}'].font = F_RED
    r += 1
    ws[f'A{r}'] = '合计（不含内部转账）'
    ws[f'C{r}'] = f'=SUMIFS(收_净额,收_归类,"工程款收款",{per(S0, S1)})'
    ws[f'D{r}'] = f'=SUMIFS(收_净额,收_类别,"收入",{gk},{per(S0, S1)})'
    ws[f'E{r}'] = f'=-SUMIFS(收_净额,收_类别,"<>收入",{gk},{per(S0, S1)})'
    ws[f'F{r}'] = f'=ROUND(C{r}+D{r}-E{r},2)'
    ws[f'G{r}'] = f'=COUNTIFS(收_有效,1,收_归类,"<>内部转账",{per(S0, S1)})'
    style_line(r, pj_cols, 'tot', pfmt, pal)
    ws[f'H{r}'] = f'=IF(ROUND(F{r}-C{net_row},2)=0,"✓ ＝②净增加","⚠ 跟②净增加差 "&ROUND(F{r}-C{net_row},2))'
    ws[f'H{r}'].font = F_NOTE
    cf_warn(ws, f'H{r}', f'H{r}')
    pj_tot = r
    r += 2

    # ═════ ④ 年报 ═════
    section(ws, r, 'A', LAST, f'="④ "&{YR}&" 年 年报（所选日期那一年；截止日以后的月份空着）"', C_VIEW)
    r += 1
    header(ws, r, [('A', '账户'), ('B', '类型'), ('C', '年初余额')] + [(VMON[m], f'{m + 1}月末余额') for m in range(12)]
           + [('P', '全年收入'), ('Q', '全年支出')], GREEN_H)
    r += 1
    y_cols = [CL(i) for i in range(1, 18)]
    yfmt = {'A': 'General', 'B': 'General'}
    yal = {'A': AL, 'B': AC}
    ysub = {}
    for g in (1, 2, 3):
        for k in range(1, M[g] + 1):
            z = hidx(r, kth(k, GCNT[g], HR0, N_ACC))
            ws[f'A{r}'] = f'=IF({z}=0,"",{hv(H["名称"], z)})'
            ws[f'B{r}'] = f'=IF({z}=0,"",{hv(H["类型"], z)})'
            ws[f'C{r}'] = f'=IF({z}=0,"",{hv(H["年初"], z)})'
            for m in range(12):
                ws[f'{VMON[m]}{r}'] = f'=IF(OR({z}=0,${MCOL[m]}$3=0),"",{hv(MCOL[m], z)})'
            ws[f'P{r}'] = f'=IF({z}=0,"",{hv(H["全年收入"], z)})'
            ws[f'Q{r}'] = f'=IF({z}=0,"",{hv(H["全年支出"], z)})'
            style_line(r, y_cols, None, yfmt, yal)
            r += 1
        ysub[g] = r
        ws[f'A{r}'] = {1: '银行＋现金 小计', 2: '专户 小计（应为 0）', 3: '个人户 小计'}[g]
        gs = lambda col, g=g: f'SUMIFS({hrng(col)},{hrng(H["组"])},{g})'
        ws[f'C{r}'] = f'={gs(H["年初"])}'
        for m in range(12):
            ws[f'{VMON[m]}{r}'] = f'=IF(${MCOL[m]}$3=0,"",{gs(MCOL[m])})'
        ws[f'P{r}'] = f'={gs(H["全年收入"])}'
        ws[f'Q{r}'] = f'={gs(H["全年支出"])}'
        style_line(r, y_cols, 'sub', yfmt, yal)
        r += 1
    ytot_acc = r
    ws[f'A{r}'] = '全部账户 合计'
    ws[f'C{r}'] = f'=SUM({hrng(H["年初"])})'
    for m in range(12):
        ws[f'{VMON[m]}{r}'] = f'=IF(${MCOL[m]}$3=0,"",SUM({hrng(MCOL[m])}))'
    ws[f'P{r}'] = f'=SUM({hrng(H["全年收入"])})'
    ws[f'Q{r}'] = f'=SUM({hrng(H["全年支出"])})'
    style_line(r, y_cols, 'tot', yfmt, yal)
    r += 1
    ws[f'A{r}'] = '其中：账户之间倒钱'
    ws[f'P{r}'] = f'={YXFER}'
    ws[f'Q{r}'] = f'={YXFER}'
    style_line(r, y_cols, None, yfmt, yal)
    for c in y_cols:
        ws[f'{c}{r}'].font = F_GRAY_I
    r += 2
    header(ws, r, [('A', '收支项目'), ('B', '归类'), ('C', '全年合计')] + [(VMON[m], f'{m + 1}月') for m in range(12)], GREEN_H)
    r += 1
    yi_cols = [CL(i) for i in range(1, 16)]
    yblocks = {}
    for side, label, cntc, nlist, mshow, nm_name, cat_name, sign, crit in (
            ('收入', '收入项目', ICNT, N_INC, M_INC, '收入项目_名称', '收入项目_归类', '', '"收入"'),
            ('支出', '支出项目', ECNT, N_EXP, M_EXP, '支出项目_名称', '支出项目_归类', '-', '"<>收入"')):
        ws[f'A{r}'] = label
        style_line(r, yi_cols, 'head', yfmt, yal)
        r += 1
        i0 = r
        for k in range(1, mshow + 1):
            z = hidx(r, kth(k, cntc, HR0, nlist))
            e = esc(f'$A{r}')
            ws[f'A{r}'] = f'=IF({z}=0,"",INDEX({nm_name},{z}))'
            ws[f'B{r}'] = f'=IF({z}=0,"",INDEX({cat_name},{z}))'
            ws[f'C{r}'] = f'=IF($A{r}="","",SUM(D{r}:O{r}))'
            for m in range(12):
                mc = MCOL[m]
                ws[f'{VMON[m]}{r}'] = (f'=IF(OR($A{r}="",${mc}$3=0),"",{sign}SUMIFS(收_净额,收_收支项目,{e},收_类别,{crit},'
                                       f'{per(f"${mc}$1", f"${mc}$2")}))')
            style_line(r, yi_cols, None, yfmt, yal)
            r += 1
        i1 = r - 1
        tot_r = r + 1
        ws[f'A{r}'] = f'其他没对上的{side}'
        ws[f'C{r}'] = f'=SUM(D{r}:O{r})'
        for m in range(12):
            c = VMON[m]
            ws[f'{c}{r}'] = f'=IF(${MCOL[m]}$3=0,"",ROUND({c}{tot_r}-SUMIFS({c}{i0}:{c}{i1},$B{i0}:$B{i1},"<>内部转账"),2))'
        style_line(r, yi_cols, None, yfmt, yal)
        r += 1
        ws[f'A{r}'] = f'{side}合计（不含内部转账）'
        ws[f'C{r}'] = f'=SUM(D{r}:O{r})'
        for m in range(12):
            mc = MCOL[m]
            ws[f'{VMON[m]}{r}'] = (f'=IF(${mc}$3=0,"",{sign}SUMIFS(收_净额,收_类别,{crit},收_归类,"<>内部转账",'
                                   f'{per(f"${mc}$1", f"${mc}$2")}))')
        style_line(r, yi_cols, 'tot', yfmt, yal)
        yblocks[side] = (i0, i1, r)
        ws.conditional_formatting.add(f'A{i0}:O{i1}', FormulaRule(formula=[f'$B{i0}="内部转账"'], font=Font(color='FF808080', italic=True)))
        r += 1
    yti, yte = yblocks['收入'][2], yblocks['支出'][2]
    ws[f'A{r}'] = '净增加（收入合计−支出合计）'
    ws[f'C{r}'] = f'=SUM(D{r}:O{r})'
    for m in range(12):
        c = VMON[m]
        ws[f'{c}{r}'] = f'=IF(${MCOL[m]}$3=0,"",ROUND({c}{yti}-{c}{yte},2))'
    style_line(r, yi_cols, 'tot', yfmt, yal)
    ynet = r
    r += 1
    put(ws, f'A{r}', '内部转账（账户间倒钱）灰色斜体，不进合计；「其他没对上的」不是 0 就去【收支登记】看校验列。', F_NOTE, align=AL, border=False)
    last_row = r

    # ── 关键结果（固定位置，首页/数据校验要用可以引用） ──
    keys = [
        (10, '本期起', f'={S0}'), (11, '本期止', f'={S1}'),
        (12, '银行＋现金期末', f'={gs_all("期末", 1)}'), (13, '专户期末', f'={gs_all("期末", 2)}'),
        (14, '个人户期末', f'={gs_all("期末", 3)}'), (15, '全部账户期末', f'=SUM({hrng(H["期末"])})'),
        (16, '本期收入合计（不含内部转账）', f'=C{ti}'), (17, '本期支出合计（不含内部转账）', f'=C{te}'),
    ]
    for rr, lbl, f in keys:
        ws[f'{SC}{rr}'] = f
        ws[f'{KEYLBL}{rr}'] = lbl
        ws[f'{SC}{rr}'].font = ws[f'{KEYLBL}{rr}'].font = F_HELP

    hide_from = CI(IDX)
    for i in range(hide_from, CI('CH') + 1):
        ws.column_dimensions[CL(i)].hidden = True
    ws.freeze_panes = 'B6'
    print_setup(ws, '3:5', landscape=True)
    ws.print_area = f'A1:{LAST}{last_row}'
    return dict(first_acc=first_acc, sub=sub, tot_acc=tot_acc, xfer=xfer_row, inc=blocks['收入'], exp=blocks['支出'],
                net=net_row, pj=(p0, p1, pj_tot), ysub=ysub, ytot=ytot_acc, yinc=yblocks['收入'], yexp=yblocks['支出'],
                ynet=ynet)


def gs_all(col, g):
    return f'SUMIFS({hrng(H[col])},{hrng(H["组"])},{g})'


# ═════════════════════════════════════════ 账户明细 ═════════════════════════════════════════
A_LAST = 'L'
A_HDR, A_R0, A_SHOW = 8, 9, 1500
A_ACC = '$AA$3'                         # 选的账户（TRIM 后；空＝全部）
A_S0, A_S1 = '$E$4', '$G$4'             # 有效起止


def build_acct(ws, ctx):
    widths(ws, {'A': 11, 'B': 16, 'C': 14, 'D': 30, 'E': 13, 'F': 13, 'G': 14, 'H': 13, 'I': 14, 'J': 16, 'K': 28, 'L': 9})
    tip = ('💡 黄格选一个账户（空＝全部账户）和起止日期（空＝首页年初～截止日），下面按日期列出这段时间的每一笔，余额逐笔往下滚——'
           '可以拿来跟银行对账单、微信账单逐笔对。别的账户转进来的钱（内部转账，「对方账户」填的是它）也列出来，收入、支出自动对调。'
           '选全部账户时，账户之间倒的钱只列一次、收入支出各记一遍，余额不变。要改哪一笔，按最后一列行号去【收支登记】改。')
    title(ws, '账户明细（逐笔带余额）', A_LAST, C_VIEW, tip)
    selector(ws, 'A3', '账户', 'B3', None, '=账户列表', prompt='从下拉选账户；空＝全部账户')
    ws.merge_cells('B3:C3')
    selector(ws, 'D3', '起', 'E3', None, fmt=DATE, prompt='空＝首页年初')
    selector(ws, 'F3', '止', 'G3', None, fmt=DATE, prompt='空＝首页截止日')
    dv_date(ws, 'E3')
    dv_date(ws, 'G3')
    home_link(ws, f'{A_LAST}3')
    ws['AA3'] = '=TRIM(B3&"")'
    ws['AA3'].font = F_HELP
    x = A_ACC
    mt = f'MATCH({x},账户_名称,0)'
    put(ws, 'A4', '实际用的', F_NOTE, align=AC)
    ws.merge_cells('B4:C4')
    cell(ws, 'B4', (f'=IF({x}="","全部账户",IF(ISNA({mt}),"⚠ 「"&{x}&"」不在【基本信息】账户表里",'
                    f'{x}&"（"&INDEX(账户_类型,{mt})&IF(INDEX(账户_所属人,{mt})<>"","·"&INDEX(账户_所属人,{mt}),"")&"）"))'),
         align=AC, font=F_AUTOB)
    put(ws, 'D4', '空＝首页年初→', F_NOTE, align=AR)
    cell(ws, 'E4', '=IF(ISNUMBER(E3),INT(E3),P_年初)', DATE, AC, F_AUTOB)
    put(ws, 'F4', '空＝首页截止日→', F_NOTE, align=AR)
    cell(ws, 'G4', '=IF(ISNUMBER(G3),INT(G3),P_截止)', DATE, AC, F_AUTOB)
    ws.merge_cells(f'H4:{A_LAST}4')
    put(ws, 'H4', f'=IF({A_S0}>{A_S1},"⚠ 起晚于止，请改日期","")', F_RED, align=AL, border=False)

    # 顶部汇总
    ws.merge_cells('A5:C6')
    put(ws, 'A5', '本期汇总（按上面的账户、起止）', F_KPI_L, fill('FFD9E1F2'), align=ACW)
    for c, t in zip('DEFGH', ('期初余额（起日前一天）', '本期收入', '本期支出', '期末余额', '笔数')):
        put(ws, f'{c}5', t, F_KPI_L, fill('FFD9E1F2'), align=ACW)
    e = esc(x)
    lt = f'"<"&{A_S0}'
    ws['D6'] = (f'=IF({x}="",ROUND(SUM(账户_期初)+SUMIFS(收_净额,收_日期,{lt})-SUMIFS(收_净额,收_归类,"内部转账",收_日期,{lt}),2),'
                f'ROUND(SUMIFS(账户_期初,账户_名称,{e})+SUMIFS(收_净额,收_账户,{e},收_日期,{lt})'
                f'-SUMIFS(收_净额,收_对方账户,{e},收_归类,"内部转账",收_日期,{lt}),2))')
    pp = per(A_S0, A_S1)
    ws['E6'] = (f'=ROUND(IF({x}="",SUMIFS(收_收入,{pp})+SUMIFS(收_支出,收_归类,"内部转账",{pp}),{acc_inc(e, A_S0, A_S1)}),2)')
    ws['F6'] = (f'=ROUND(IF({x}="",SUMIFS(收_支出,{pp})+SUMIFS(收_收入,收_归类,"内部转账",{pp}),{acc_exp(e, A_S0, A_S1)}),2)')
    ws['G6'] = '=ROUND(D6+E6-F6,2)'
    for c in 'DEFG':
        put(ws, f'{c}6', None, F_KPI_V, FILL_AUTO, MONEY, AR)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 24

    # 排序键（覆盖 _收 全部 N_CASH 条）
    def cond(i):
        n = i + 1
        return (f'AND(INDEX(收_有效,{n})=1,INDEX(收_日期,{n})>={A_S0},INDEX(收_日期,{n})<={A_S1},'
                f'OR({x}="",INDEX(收_账户,{n})={x},AND(INDEX(收_对方账户,{n})={x},INDEX(收_归类,{n})="内部转账")))')
    ncell = skey(ws, 'Z', A_R0, N_CASH, cond, lambda i: f'INDEX(收_日期,{i + 1})')
    put(ws, 'H6', f'={ncell}', F_KPI_V, FILL_AUTO, INT, AC)
    last_bal = f'INDEX($G${A_R0}:$G${A_R0 + A_SHOW - 1},{ncell})'
    ws.merge_cells(f'A7:{A_LAST}7')
    put(ws, 'A7', (f'=IF({A_S0}>{A_S1},"⚠ 起晚于止，请改日期",IF({ncell}=0,"这段时间没有流水",'
                   f'IF({ncell}>{A_SHOW},"⚠ 共 "&{ncell}&" 笔，只显示前 {A_SHOW} 笔（余额也只滚到第 {A_SHOW} 笔），请缩短日期范围",'
                   f'"共 "&{ncell}&" 笔，按日期排（同一天按录入顺序）"&IF(ROUND({last_bal}-$G$6,2)=0,"；最后一笔的余额＝期末余额 ✓",'
                   f'"；⚠ 最后一笔的余额跟期末余额对不上"))))'), F_NOTE, align=AL, border=False)
    cf_warn(ws, 'A7', 'A7')

    heads = [('A', '日期'), ('B', '账户'), ('C', '收支项目'), ('D', '摘要'), ('E', '收入'), ('F', '支出'), ('G', '余额'),
             ('H', '项目名称'), ('I', '客户/供应商/人员'), ('J', '对方账户 / 转出账户'), ('K', '校验'), ('L', '收支登记\n行号')]
    header(ws, A_HDR, heads, GREEN_H)
    fm = {'A': DATE, 'E': MONEY_B, 'F': MONEY_B, 'G': MONEY, 'L': '0'}
    al = {'A': AC, 'B': AL, 'C': AL, 'D': AL, 'E': AR, 'F': AR, 'G': AR, 'H': AL, 'I': AL, 'J': AL, 'K': AL, 'L': AC}
    for i in range(A_SHOW):
        r = A_R0 + i
        y = f'$Y{r}'
        ws[f'X{r}'] = i + 1
        ws[f'Y{r}'] = '=' + ksorted(f'$X{r}', 'Z', A_R0, N_CASH, ncell)
        ws[f'W{r}'] = f'=IF({y}=0,0,IF(AND({x}<>"",INDEX(收_账户,{y})<>{x}),1,0))'
        ws[f'V{r}'] = f'=IF({y}=0,0,IF(AND({x}="",INDEX(收_归类,{y})="内部转账"),1,0))'
        g = lambda f: f'INDEX(收_{f},{y})'
        w, v = f'$W{r}', f'$V{r}'
        both = f'{g("收入")}+{g("支出")}'
        f = {
            'A': f'=IF({y}=0,"",{g("日期")})',
            'B': (f'=IF({y}=0,"",IF({w}=1,{x},IF({v}=1,IF({g("净额")}<0,{g("账户")}&"→"&{g("对方账户")},'
                  f'{g("对方账户")}&"→"&{g("账户")}),{g("账户")})))'),
            'C': f'=IF({y}=0,"",{g("收支项目")})',
            'D': f'=IF({y}=0,"",{g("摘要")})',
            'E': f'=IF({y}=0,"",IF({w}=1,{g("支出")},IF({v}=1,{both},{g("收入")})))',
            'F': f'=IF({y}=0,"",IF({w}=1,{g("收入")},IF({v}=1,{both},{g("支出")})))',
            'G': f'=IF({y}=0,"",ROUND({"$D$6" if i == 0 else f"G{r - 1}"}+E{r}-F{r},2))',
            'H': f'=IF({y}=0,"",{g("项目原")})',
            'I': f'=IF({y}=0,"",{g("往来对象")})',
            'J': f'=IF({y}=0,"",IF({w}=1,{g("账户")},IF({v}=1,"（账户间倒钱，两边抵消）",{g("对方账户")})))',
            'K': f'=IF({y}=0,"",{g("校验")})',
            'L': f'=IF({y}=0,"",{g("录入行")})',
        }
        for c, val in f.items():
            q = ws[f'{c}{r}']
            q.value = val
            q.font = F_TXT
            q.number_format = fm.get(c, 'General')
            q.alignment = al[c]
        for c in 'VWXY':
            ws[f'{c}{r}'].font = F_HELP
    r1 = A_R0 + A_SHOW - 1
    rng = f'A{A_R0}:{A_LAST}{r1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$Y{A_R0}>0'], border=BD))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'$C{A_R0}="内部转账"'], fill=fill('FFDDEBF7')))
    ws.conditional_formatting.add(f'K{A_R0}:K{r1}', FormulaRule(formula=[f'LEFT($K{A_R0},1)="✗"'], font=F_RED_CF))
    ws.conditional_formatting.add(f'K{A_R0}:K{r1}', FormulaRule(formula=[f'LEFT($K{A_R0},1)="⚠"'], font=Font(color='FFBF8F00')))
    for i in range(CI('V'), CI('AA') + 1):
        ws.column_dimensions[CL(i)].hidden = True
    ws.freeze_panes = f'B{A_R0}'
    print_setup(ws, f'{A_HDR}:{A_HDR}', landscape=True)
    ws.print_area = f'A1:{A_LAST}{r1}'


def build(wb, ctx):
    ctx = ctx or {}
    build_fund(wb[SH_FUND], ctx)
    build_acct(wb[SH_ACCT], ctx)
