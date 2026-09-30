# -*- coding: utf-8 -*-
"""把各张表拼成一本：录入（蓝）→ 档案（灰）→ 查看（绿）→ 项目（橙）→ 老板报表（深红）→ 数据校验；隐藏的取数表放最后。"""
import openpyxl
from openpyxl.styles import Protection
from openpyxl.worksheet.protection import SheetProtection
from common import home_link, print_setup, CL
from layout import *
import demo
import s_base, s_master, s_open, s_cash, s_regs, s_att, s_agg

TAB = {'录入': 'FF2F75B5', '档案': 'FF7F7F7F', '查看': 'FF548235', '项目': 'FFC65911', '报表': 'FF833C0C', '首页': 'FF1F3864'}
GROUPS = [
    ('首页', [SH_HOME]),
    ('录入', [SH_CASH, SH_AP, SH_ATT, SH_INV, SH_REV, SH_OFF]),
    ('档案', [SH_PROJ, SH_UNIT, SH_RATE, SH_OPEN, SH_BASE]),
    ('查看', [SH_ACC, '资金报表', SH_AR, SH_APS, SH_PER, SH_PAY, SH_PAYS, SH_LAB, SH_INVS]),
    ('项目', [SH_PL, SH_PPL, SH_ALLOC]),
    ('报表', [SH_IS, SH_BS, SH_BE, SH_CHK]),
]
HIDDEN = [SH_MS, SH_PS, SH_BALX, SH_AUX]
# 打印：表头行（每页重复）、横向/竖向
PRINT = {SH_CASH: ('5:5', True), SH_ATT: ('4:4', True), SH_AP: ('4:4', True), SH_INV: ('4:4', True), SH_ACC: ('7:7', True),
         SH_AR: ('5:5', True), SH_APS: ('6:6', True), SH_PER: ('5:5', True), SH_PAYS: ('5:5', True), SH_LAB: ('5:5', True),
         SH_INVS: ('5:5', True), SH_PL: (None, True), SH_PPL: ('4:5', True), SH_ALLOC: (None, True), SH_IS: ('5:5', True),
         SH_BS: ('5:5', False), SH_BE: (None, False), SH_CHK: ('4:4', False), SH_HOME: (None, False)}
KEEP_OPEN = {'FFFFFF00', 'FFFFF7E0', 'FFEAF1FB'}      # 黄色选择格、淡黄手填格、淡蓝粘贴格：保护时不锁


def _balanced(f):
    depth, q = 0, False
    for ch in f:
        if ch == '"':
            q = not q
        elif not q:
            depth += (ch == '(') - (ch == ')')
            if depth < 0:
                return False
    return depth == 0 and not q


def _check_formulas(wb):
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    if '_xlfn.' in v or not _balanced(v):
                        bad.append(f'{ws.title}!{c.coordinate}: {v[:120]}')
    assert not bad, '\n'.join(bad[:20])


def build(scenario=None):
    ctx = demo.build_ctx()
    if scenario:
        scenario(ctx)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    s_cash.build_cash(wb, ctx)
    s_regs.build_ap(wb, ctx)
    s_att.build_att(wb, ctx)
    s_regs.build_inv(wb, ctx)
    s_regs.build_rev(wb, ctx)
    s_regs.build_off(wb, ctx)
    s_master.build_proj(wb, ctx)
    s_master.build_unit(wb, ctx)
    s_master.build_rate(wb, ctx)
    s_open.build_open(wb, ctx)
    s_base.build_base(wb, ctx)
    s_agg.build_ms(wb, ctx)
    s_agg.build_ps(wb, ctx)
    s_agg.build_balx(wb, ctx)
    import s_views, s_views2, s_proj, s_fin, s_check
    for m in (s_views, s_views2, s_proj, s_fin, s_check):
        m.build_all(wb, ctx)
    s_base.build_aux(wb, ctx)
    order = []
    for grp, names in GROUPS:
        for n in names:
            if n in wb.sheetnames:
                order.append(n)
                wb[n].sheet_properties.tabColor = TAB[grp][2:]
    order += [n for n in HIDDEN if n in wb.sheetnames]
    rest = [n for n in wb.sheetnames if n not in order]
    assert not rest, rest
    wb._sheets = [wb[n] for n in order]
    for grp, names in GROUPS:
        for n in names:
            if n not in wb.sheetnames:
                continue
            ws = wb[n]
            if grp in ('录入', '档案'):
                end = max((m.max_col for m in ws.merged_cells.ranges if m.min_row == 1 and m.min_col == 1), default=ws.max_column)
                home_link(ws, f'{CL(end + 1)}1')
                ws.column_dimensions[CL(end + 1)].width = max(ws.column_dimensions[CL(end + 1)].width or 0, 10)
            if grp in ('首页', '查看', '项目', '报表'):
                for row in ws.iter_rows():
                    for c in row:
                        if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb in KEEP_OPEN:
                            c.protection = Protection(locked=False)
                ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                                formatCells=False, selectLockedCells=False, selectUnlockedCells=False)
            if n in PRINT:
                rows_, land = PRINT[n]
                print_setup(ws, rows_, landscape=land)
    _check_formulas(wb)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 0
    wb.worksheets[0].sheet_view.tabSelected = True
    wb.calculation.fullCalcOnLoad = False
    wb.calculation.calcId = 191029
    return wb
