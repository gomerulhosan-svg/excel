# -*- coding: utf-8 -*-
"""【发票统计】：全公司开出去的票（销项）、收到的票（进项）按项目、按供应商、按月统计，估算增值税，跟税局导出（【发票登记】）逐月核对。
   ① 销项按项目（应收登记 类型＝开票）；② 进项按供应商（应付登记 已开票）＋ 收支登记零星拿到的票；③ 按月：销项/进项/估算增值税/附加；
   ④ 按月跟税局导出核对；⑤ 税局导出按对方汇总。起止黄格在 B3/B4（空＝首页年初/截止），实际用的在 C3/C4。
   清单用本表右边隐藏列（BA～BE 计数、CA～ 条件）只列有数的行；合计行在表头下面第一行（不受清单显示行数限制，直接对整块 SUMIFS）。"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *          # 放在 common 后面：颜色常量以 layout 为准

LASTC = 'N'
IN_START, IN_END = 'B3', 'B4'             # 黄格：起、止
S, E = '$C$3', '$C$4'                     # 实际用的起、止
CAP_PJ, CAP_SUP, CAP_MON, CAP_CP = 25, 50, 24, 200
H0 = 5                                    # 隐藏辅助列从第 5 行起
HC_PJ, HC_SUP, HC_XS, HC_JX, HC_KEY = 'BA', 'BB', 'BC', 'BD', 'BE'
IDX, AUX, AUX2 = 'P', 'Q', 'R'            # 显示行的隐藏列：P 第几条（③④＝这行的起）、Q（③④＝这行的止；⑤＝方向|对方）、R（⑤＝对方）
C_SEC = 'FF375623'                        # 区块标题条（深绿）
RATES_OUT = [('9%', 0.09), ('13%', 0.13), ('6%', 0.06), ('3%', 0.03), ('1%', 0.01)]
RATES_IN = [('13%', 0.13), ('9%', 0.09), ('6%', 0.06), ('3%', 0.03), ('1%', 0.01)]
FMT_MON = 'yyyy"年"m"月"'
F_KV = Font(name=YH, sz=11, bold=True, color='FFC00000')

# ───────── 行号（固定，首页/数据校验可以引用） ─────────
R_PJ = 6                                   # ① 区块标题
R_PJ_H, R_PJ_T = R_PJ + 1, R_PJ + 2        # 表头、合计
R_PJ_0 = R_PJ + 3                          # 第 1 个项目
R_SUP = R_PJ_0 + CAP_PJ + 1                # ②
R_SUP_H, R_SUP_T, R_SUP_X, R_SUP_C = R_SUP + 1, R_SUP + 2, R_SUP + 3, R_SUP + 4   # 表头、合计、其中名单外、收支登记零星票
R_SUP_0 = R_SUP + 5
R_MON = R_SUP_0 + CAP_SUP + 1              # ③
R_MON_H, R_MON_T = R_MON + 1, R_MON + 2
R_MON_0 = R_MON + 3
R_CHK = R_MON_0 + CAP_MON + 1              # ④
R_CHK_N, R_CHK_H, R_CHK_T = R_CHK + 1, R_CHK + 2, R_CHK + 3   # 说明行（隐藏 P 列放税局有效张数）、表头、合计
R_CHK_0 = R_CHK + 4
R_CP = R_CHK_0 + CAP_MON + 1               # ⑤
R_CP_N, R_CP_H, R_CP_TS, R_CP_TJ = R_CP + 1, R_CP + 2, R_CP + 3, R_CP + 4   # 说明行、表头、销项合计、进项合计
R_CP_0 = R_CP + 5
R_LAST = R_CP_0 + CAP_CP - 1
PIAO_N = f'${IDX}${R_CHK_N}'               # 发票登记里有效的票（全部日期）张数

# 关键结果格（首页、数据校验可引用）
KPI = dict(销项开票='I3', 销项税='I4', 收到的票='K3', 可抵进项税='K4', 估算应交增值税='M3', 跟税局差额='M4',
           税局销项差额=f'D{R_CHK_T}', 税局进项差额=f'H{R_CHK_T}')


def dr(rng, lo=S, hi=E):
    """日期段条件"""
    return f'{rng},">="&{lo},{rng},"<="&{hi}'


def rate_crit(rng, r):
    """税率条件：r±0.5%（手填的 9%、0.09 都算 9%）"""
    return f'{rng},">{round(r - 0.005, 4)}",{rng},"<{round(r + 0.005, 4)}"'


def cash_inv(lo=S, hi=E):
    """收支登记里拿到的票：按开票日期；开票日期没填的按付款日期"""
    return (f'SUMIFS(收_已开票,{dr("收_开票日期", lo, hi)})'
            f'+SUMIFS(收_已开票,收_开票日期,0,{dr("收_日期", lo, hi)})')


def _cells(ws, r, cols, fmt=MONEY, font=F_TXT, fill_=None, border=False, align=AR):
    for c in cols:
        put(ws, f'{c}{r}', None, font, fill_, fmt, align, border=border)


def _total_row(ws, r, c1, c2):
    for i in range(CI(c1), CI(c2) + 1):
        c = ws.cell(row=r, column=i)
        c.font, c.fill, c.border = F_TXTB, FILL_TOT, BD
        if c.number_format == 'General':
            c.number_format = MONEY
        c.alignment = AR if i > 1 else ALW
    ws.row_dimensions[r].height = 30


def _list_style(ws, r0, n, c2, fmts, left='A', center='', key='A'):
    """清单行：平时不画框，有内容的行用条件格式画框（不显示一大片空格子）"""
    for r in range(r0, r0 + n):
        for i in range(1, CI(c2) + 1):
            col = CL(i)
            al = AL if col in left else (AC if col in center else AR)
            put(ws, f'{col}{r}', None, F_TXT, None, fmts.get(col, MONEY), al, border=False)
    ws.conditional_formatting.add(f'A{r0}:{c2}{r0 + n - 1}', FormulaRule(formula=[f'${key}{r0}<>""'], border=BD))


# ───────────────────────── ① 销项按项目 ─────────────────────────
def blk_out(ws):
    section(ws, R_PJ, 'A', LASTC, '① 开出去的票（销项）· 按项目　——　来自【应收登记】类型「开票」，按开票日期', C_SEC)
    heads = [('A', '项目'), ('B', '甲方'), ('C', '开票金额\n（含税）')]
    heads += [(CL(4 + i), f'{t}\n税率') for i, (t, _r) in enumerate(RATES_OUT)]
    heads += [('I', '其他税率'), ('J', '销项税'), ('K', '张数')]
    header(ws, R_PJ_H, heads, C_VIEW)
    # 隐藏：哪些项目这段时间开过票
    nm = lambda i: f'INDEX(项目_名称,{i + 1})'
    counter(ws, HC_PJ, H0, N_PJ,
            lambda i: f'IF({nm(i)}="",FALSE,COUNTIFS(应_项目,{esc(nm(i))},应_开票额,"<>0",{dr("应_日期")})>0)')
    n = cnt(HC_PJ, H0, N_PJ)

    def row(r, pj):
        cond = f'应_项目,{esc(pj)},' if pj else ''
        g = (lambda f: f'=IF($A{r}="","",{f})') if pj else (lambda f: f'={f}')
        ws[f'C{r}'] = g(f'SUMIFS(应_开票额,{cond}{dr("应_日期")})')
        for i, (_t, rt) in enumerate(RATES_OUT):
            ws[f'{CL(4 + i)}{r}'] = g(f'SUMIFS(应_开票额,{cond}{rate_crit("应_税率", rt)},{dr("应_日期")})')
        ws[f'I{r}'] = g(f'ROUND(C{r}-SUM(D{r}:H{r}),2)')
        ws[f'J{r}'] = g(f'SUMIFS(应_销项税,{cond}{dr("应_日期")})')
        ws[f'K{r}'] = g(f'COUNTIFS({cond}应_开票额,"<>0",{dr("应_日期")})')

    r = R_PJ_T
    ws[f'A{r}'] = f'="合计（共 "&{n}&" 个项目开了票"&IF({n}>{CAP_PJ},"，下面只列前 {CAP_PJ} 个","")&"）"'
    row(r, None)
    _total_row(ws, r, 'A', 'K')
    ws.merge_cells(f'A{r}:B{r}')
    ws[f'K{r}'].number_format = INT
    fm = {'K': INT}
    _list_style(ws, R_PJ_0, CAP_PJ, 'K', fm, left='AB')
    for k in range(1, CAP_PJ + 1):
        r = R_PJ_0 + k - 1
        ws[f'{IDX}{r}'] = f'={kth(k, HC_PJ, H0, N_PJ)}'
        ws[f'A{r}'] = f'=IF(${IDX}{r}=0,"",INDEX(项目_名称,${IDX}{r}))'
        ws[f'B{r}'] = f'=IF($A{r}="","",INDEX(项目_甲方,${IDX}{r}))'
        row(r, f'$A{r}')


# ───────────────────────── ② 进项按供应商 ─────────────────────────
def blk_in(ws):
    section(ws, R_SUP, 'A', LASTC, '② 收到的票（进项）· 按供应商　——　来自【应付登记】已开票金额（按开票日期；没填开票日期按那行日期），'
                                   '应付额按应付日期', C_SEC)
    heads = [('A', '供应商'), ('B', '类型'), ('C', '收票合计\n（含税）'), ('D', '专票'), ('E', '普票\n（含没写类型的）')]
    heads += [(CL(6 + i), f'{t}\n税率') for i, (t, _r) in enumerate(RATES_IN)]
    heads += [('K', '其他税率'), ('L', '可抵进项税\n（专票）'), ('M', '应付额\n（同期）'), ('N', '未开票\n（应付−收票）')]
    header(ws, R_SUP_H, heads, C_VIEW)
    nm = lambda i: f'INDEX(供应商_名称,{i + 1})'
    counter(ws, HC_SUP, H0, N_SUP,
            lambda i: (f'IF({nm(i)}="",FALSE,COUNTIFS(付_供应商,{esc(nm(i))},付_已开票,"<>0",{dr("付_开票日期")})'
                       f'+COUNTIFS(付_供应商,{esc(nm(i))},付_应付额,"<>0",{dr("付_日期")})>0)'))
    n = cnt(HC_SUP, H0, N_SUP)
    l0, l1 = R_SUP_0, R_SUP_0 + CAP_SUP - 1

    def row(r, sup):
        cond = f'付_供应商,{esc(sup)},' if sup else ''
        g = (lambda f: f'=IF($A{r}="","",{f})') if sup else (lambda f: f'={f}')
        ws[f'C{r}'] = g(f'SUMIFS(付_已开票,{cond}{dr("付_开票日期")})')
        ws[f'D{r}'] = g(f'SUMIFS(付_已开票,{cond}付_发票类型,"专票",{dr("付_开票日期")})')
        ws[f'E{r}'] = g(f'ROUND(C{r}-D{r},2)')
        for i, (_t, rt) in enumerate(RATES_IN):
            ws[f'{CL(6 + i)}{r}'] = g(f'SUMIFS(付_已开票,{cond}{rate_crit("付_税率", rt)},{dr("付_开票日期")})')
        ws[f'K{r}'] = g(f'ROUND(C{r}-SUM(F{r}:J{r}),2)')
        ws[f'L{r}'] = g(f'SUMIFS(付_进项税,{cond}{dr("付_开票日期")})')
        ws[f'M{r}'] = g(f'SUMIFS(付_应付额,{cond}{dr("付_日期")})')
        ws[f'N{r}'] = g(f'ROUND(M{r}-C{r},2)')

    r = R_SUP_T
    ws[f'A{r}'] = f'="合计（应付登记，共 "&{n}&" 家"&IF({n}>{CAP_SUP},"，下面只列前 {CAP_SUP} 家","")&"）"'
    row(r, None)
    _total_row(ws, r, 'A', 'N')
    ws.merge_cells(f'A{r}:B{r}')
    # 其中：没列出来的（供应商不在【供应商信息】里、或超过显示行数）——只有不为 0 才显示
    rx, t = R_SUP_X, R_SUP_T
    res = lambda c: f'ROUND({c}${t}-SUM({c}${l0}:{c}${l1}),2)'
    ws[f'A{rx}'] = (f'=IF(OR(ABS({res("C")})>=0.01,ABS({res("M")})>=0.01),'
                    f'"其中：下面没列出来的（供应商不在【供应商信息】里，或超过 {CAP_SUP} 家）","")')
    for c in 'CDEFGHIJKLMN':
        ws[f'{c}{rx}'] = f'=IF($A{rx}="","",{res(c)})'
    _cells(ws, rx, 'CDEFGHIJKLMN')
    put(ws, f'A{rx}', None, F_NOTE, None, None, ALW, border=False)
    ws.merge_cells(f'A{rx}:B{rx}')
    ws.conditional_formatting.add(f'A{rx}:{LASTC}{rx}', FormulaRule(formula=[f'$A{rx}<>""'], border=BD))
    # 收支登记里零星拿到的票
    rc = R_SUP_C
    ws[f'A{rc}'] = '另：【收支登记】里零星采购、垫付拿到的票（没有税率，不分专票普票）'
    ws[f'C{rc}'] = f'={cash_inv()}'
    for c in 'ABCDEFGHIJKLMN':
        put(ws, f'{c}{rc}', None, F_TXTB if c in 'AC' else F_TXT, FILL_SUB, MONEY, ALW if c == 'A' else AR)
    ws.merge_cells(f'A{rc}:B{rc}')
    ws.row_dimensions[rc].height = 30
    _list_style(ws, l0, CAP_SUP, LASTC, {}, left='A', center='B')
    for k in range(1, CAP_SUP + 1):
        r = l0 + k - 1
        ws[f'{IDX}{r}'] = f'={kth(k, HC_SUP, H0, N_SUP)}'
        ws[f'A{r}'] = f'=IF(${IDX}{r}=0,"",INDEX(供应商_名称,${IDX}{r}))'
        ws[f'B{r}'] = f'=IF($A{r}="","",INDEX(供应商_类型,${IDX}{r}))'
        row(r, f'$A{r}')


# ───────────────────────── 月份行（③④共用写法） ─────────────────────────
def _month_cols(ws, r, k):
    """P＝这一行的起（月初，第一个月是「起」），Q＝这一行的止（月末，最后一个月是「止」）；超出「止」的行为空"""
    ms = f'DATE(YEAR({S}),MONTH({S})+{k - 1},1)'
    me = f'DATE(YEAR({S}),MONTH({S})+{k},0)'
    ws[f'{IDX}{r}'] = f'=IF({ms}>{E},"",MAX({S},{ms}))'
    ws[f'{AUX}{r}'] = f'=IF(${IDX}{r}="","",MIN({E},{me}))'
    ws[f'A{r}'] = f'=IF(${IDX}{r}="","",${IDX}{r})'
    return f'${IDX}{r}', f'${AUX}{r}'


NMON = f'((YEAR({E})-YEAR({S}))*12+MONTH({E})-MONTH({S})+1)'


# ───────────────────────── ③ 按月 ─────────────────────────
def blk_month(ws):
    section(ws, R_MON, 'A', LASTC, '③ 按月：开票、收票、估算增值税　——　选的起止之间逐月（最多 24 个月）；负数＝进项比销项多（留抵）', C_SEC)
    heads = [('A', '月份'), ('B', '销项开票额\n（含税）'), ('C', '销项税'), ('D', '进项专票额\n（含税）'), ('E', '可抵进项税'),
             ('F', '估算应交增值税\n（销项税−进项税）'), ('G', '附加税\n（×附加税率）'), ('H', '税金合计\n（增值税＋附加）'),
             ('I', '收支登记\n零星票（没税率）')]
    header(ws, R_MON_H, heads, C_VIEW)

    def row(r, lo, hi, guard):
        g = (lambda f: f'=IF({guard}="","",{f})') if guard else (lambda f: f'={f}')
        ws[f'B{r}'] = g(f'SUMIFS(应_开票额,{dr("应_日期", lo, hi)})')
        ws[f'C{r}'] = g(f'SUMIFS(应_销项税,{dr("应_日期", lo, hi)})')
        ws[f'D{r}'] = g(f'SUMIFS(付_已开票,付_发票类型,"专票",{dr("付_开票日期", lo, hi)})')
        ws[f'E{r}'] = g(f'SUMIFS(付_进项税,{dr("付_开票日期", lo, hi)})')
        ws[f'F{r}'] = g(f'ROUND(C{r}-E{r},2)')
        ws[f'G{r}'] = g(f'ROUND(F{r}*P_附加税率,2)')
        ws[f'H{r}'] = g(f'ROUND(F{r}+G{r},2)')
        ws[f'I{r}'] = g(cash_inv(lo, hi))

    r = R_MON_T
    ws[f'A{r}'] = f'=IF({NMON}>{CAP_MON},"合计（"&{NMON}&" 个月，下面只列前 {CAP_MON} 个月）","合计")'
    row(r, S, E, None)
    _total_row(ws, r, 'A', 'I')
    _list_style(ws, R_MON_0, CAP_MON, 'I', {'A': FMT_MON}, left='', center='A')
    for k in range(1, CAP_MON + 1):
        r = R_MON_0 + k - 1
        lo, hi = _month_cols(ws, r, k)
        row(r, lo, hi, lo)


# ───────────────────────── ④ 跟税局导出核对 ─────────────────────────
def blk_check(ws):
    section(ws, R_CHK, 'A', LASTC, '④ 跟税局导出的发票核对（按月）　——　税局＝【发票登记】（作废的不算），差额＝税局−我们登记的', C_SEC)
    rn = R_CHK_N
    ws[PIAO_N.replace('$', '')] = '=COUNTIF(票_有效,1)'
    ws.merge_cells(f'A{rn}:{LASTC}{rn}')
    put(ws, f'A{rn}', (f'=IF({PIAO_N}=0,"还没粘贴税局导出的发票——到【发票登记】把电子税务局导出的「发票基础信息」粘进去，这里才能逐月核对。",'
                       f'IF(P_税号="","⚠ 【基本信息】① 没填公司税号：发票登记只能按公司名称分开出去/收到的票，名称跟税局写的不一样就会分错。",'
                       f'"税局导出按开票日期、价税合计算；差额不为 0 的月份，看最右边的提示去找是哪张票。"))'),
        F_RED, None, None, ALW, border=False)
    ws.row_dimensions[rn].height = 20
    heads = [('A', '月份'), ('B', '税局·销项\n（价税合计）'), ('C', '应收登记\n开票额'), ('D', '销项差额'),
             ('E', '税局·进项\n（价税合计）'), ('F', '应付登记\n已开票'), ('G', '收支登记\n已开票'), ('H', '进项差额'), ('I', '提示')]
    header(ws, R_CHK_H, heads, C_VIEW)
    ws.merge_cells(f'I{R_CHK_H}:{LASTC}{R_CHK_H}')
    msg = '去【发票登记】和【应收/应付登记】对一下哪张票没录'

    def row(r, lo, hi, guard, total=False):
        g = (lambda f: f'=IF({guard}="","",{f})') if guard else (lambda f: f'={f}')
        gp = (lambda f: f'=IF(OR({guard}="",{PIAO_N}=0),"",{f})') if guard else (lambda f: f'=IF({PIAO_N}=0,"",{f})')
        ws[f'B{r}'] = gp(f'SUMIFS(票_价税合计,票_方向,"销项",{dr("票_日期", lo, hi)})')
        ws[f'C{r}'] = g(f'SUMIFS(应_开票额,{dr("应_日期", lo, hi)})')
        ws[f'D{r}'] = f'=IF(B{r}="","",ROUND(B{r}-C{r},2))'
        ws[f'E{r}'] = gp(f'SUMIFS(票_价税合计,票_方向,"进项",{dr("票_日期", lo, hi)})')
        ws[f'F{r}'] = g(f'SUMIFS(付_已开票,{dr("付_开票日期", lo, hi)})')
        ws[f'G{r}'] = g(cash_inv(lo, hi))
        ws[f'H{r}'] = f'=IF(E{r}="","",ROUND(E{r}-F{r}-G{r},2))'
        bad = (f'IF(AND(ABS(D{r})>=0.01,ABS(H{r})>=0.01),"销项、进项都对不上：",'
               f'IF(ABS(D{r})>=0.01,"销项对不上：","进项对不上："))')
        ok = '"✓ 整个期间都对得上"' if total else '"✓ 一致"'
        tip = (f'IF({PIAO_N}=0,"还没粘贴税局导出的发票",IF(AND(ABS(D{r})<0.01,ABS(H{r})<0.01),{ok},'
               f'{bad}&"{msg}"))')
        ws[f'I{r}'] = f'=IF($A{r}="","",{tip})' if guard else f'={tip}'
        ws.merge_cells(f'I{r}:{LASTC}{r}')

    r = R_CHK_T
    ws[f'A{r}'] = f'=IF({NMON}>{CAP_MON},"合计（"&{NMON}&" 个月，下面只列前 {CAP_MON} 个月）","合计")'
    row(r, S, E, None, total=True)
    _total_row(ws, r, 'A', LASTC)
    ws[f'I{r}'].alignment = ALW
    _list_style(ws, R_CHK_0, CAP_MON, LASTC, {'A': FMT_MON}, left='I', center='A')
    for k in range(1, CAP_MON + 1):
        r = R_CHK_0 + k - 1
        lo, hi = _month_cols(ws, r, k)
        row(r, lo, hi, lo)
    # 差额不为 0 的格子标红
    ws.conditional_formatting.add(f'D{R_CHK_T}:D{R_CHK_0 + CAP_MON - 1}',
                                  FormulaRule(formula=[f'AND(ISNUMBER(D{R_CHK_T}),ABS(N(D{R_CHK_T}))>=0.01)'], fill=FILL_WARN))
    ws.conditional_formatting.add(f'H{R_CHK_T}:H{R_CHK_0 + CAP_MON - 1}',
                                  FormulaRule(formula=[f'AND(ISNUMBER(H{R_CHK_T}),ABS(N(H{R_CHK_T}))>=0.01)'], fill=FILL_WARN))


# ───────────────────────── ⑤ 税局导出按对方 ─────────────────────────
def blk_cp(ws):
    section(ws, R_CP, 'A', LASTC, '⑤ 税局导出的发票 · 按对方汇总　——　【发票登记】选的起止之间（按开票日期，作废的不算）', C_SEC)
    rn = R_CP_N
    ws.merge_cells(f'A{rn}:{LASTC}{rn}')
    put(ws, f'A{rn}', (f'=IF({PIAO_N}=0,"还没粘贴税局导出的发票。",'
                       f'"对方＝开出去的票看购方、收到的票看销方（税局写的全称，可能跟【供应商信息】里的简称不一样）。'
                       f'税局导出的「发票基础信息」是按整张票汇总的、没有税率列，这里的税率＝税额÷金额（一张票几种税率时是平均数）。")'),
        F_NOTE, None, None, ALW, border=False)
    heads = [('A', '对方（税局写的名称）'), ('B', '方向'), ('C', '张数'), ('D', '金额\n（不含税）'), ('E', '税额'), ('F', '价税合计'),
             ('G', '税率\n（税额÷金额）')]
    header(ws, R_CP_H, heads, C_VIEW)
    # 隐藏：第 i 张票的「方向|对方」（不在起止内、无效＝空）；销项、进项各自第一次出现的计数
    n1 = N_INV
    for i in range(n1):
        r = H0 + i
        ws[f'{HC_KEY}{r}'] = (f'=IF(INDEX(票_有效,{i + 1})<>1,"",IF(INDEX(票_日期,{i + 1})<{S},"",IF(INDEX(票_日期,{i + 1})>{E},"",'
                              f'INDEX(票_方向,{i + 1})&"|"&INDEX(票_对方,{i + 1}))))')
        ws[f'{HC_KEY}{r}'].font = F_HELP
    key = lambda i: f'${HC_KEY}{H0 + i}'

    def first(direction):
        def cond(i):
            if i == 0:
                return f'IF({key(i)}="",FALSE,LEFT({key(i)},2)="{direction}")'
            return (f'IF({key(i)}="",FALSE,IF(LEFT({key(i)},2)<>"{direction}",FALSE,'
                    f'COUNTIF(${HC_KEY}${H0}:${HC_KEY}{H0 + i - 1},{esc(key(i))})=0))')
        return cond

    counter(ws, HC_XS, H0, n1, first('销项'))
    counter(ws, HC_JX, H0, n1, first('进项'))
    ns, nj = cnt(HC_XS, H0, n1), cnt(HC_JX, H0, n1)
    keys = f'${HC_KEY}${H0}:${HC_KEY}${H0 + n1 - 1}'

    def row(r, name, direction):
        cond = f'票_对方,{esc(name)},' if name else ''
        g = (lambda f: f'=IF(${AUX}{r}="","",{f})') if name else (lambda f: f'={f}')
        ws[f'C{r}'] = g(f'COUNTIFS({cond}票_方向,{direction},票_有效,1,{dr("票_日期")})')
        ws[f'D{r}'] = g(f'SUMIFS(票_金额,{cond}票_方向,{direction},{dr("票_日期")})')
        ws[f'E{r}'] = g(f'SUMIFS(票_税额,{cond}票_方向,{direction},{dr("票_日期")})')
        ws[f'F{r}'] = g(f'SUMIFS(票_价税合计,{cond}票_方向,{direction},{dr("票_日期")})')
        ws[f'G{r}'] = g(f'IF(D{r}=0,0,ROUND(E{r}/D{r},4))')

    for r, d, lbl, nn in ((R_CP_TS, '销项', '销项合计（开出去的票，', ns), (R_CP_TJ, '进项', '进项合计（收到的票，', nj)):
        tail = f'&IF({ns}+{nj}>{CAP_CP},"；下面只列前 {CAP_CP} 个对方","")' if d == '进项' else ''
        ws[f'A{r}'] = f'="{lbl}"&{nn}&" 个对方）"{tail}'
        ws[f'B{r}'] = d
        row(r, None, f'"{d}"')
        _total_row(ws, r, 'A', 'G')
        ws[f'B{r}'].alignment = AC
        ws[f'C{r}'].number_format = INT
        ws[f'G{r}'].number_format = PCT
    _list_style(ws, R_CP_0, CAP_CP, 'G', {'C': INT, 'G': PCT}, left='A', center='B')
    for k in range(1, CAP_CP + 1):
        r = R_CP_0 + k - 1
        ws[f'{IDX}{r}'] = (f'=IF({k}<={ns},{kth(k, HC_XS, H0, n1)},{kth(f"({k}-{ns})", HC_JX, H0, n1)})')
        ws[f'{AUX}{r}'] = f'=IF(${IDX}{r}=0,"",INDEX({keys},${IDX}{r}))'
        ws[f'{AUX2}{r}'] = f'=IF(${AUX}{r}="","",MID(${AUX}{r},4,200))'
        ws[f'A{r}'] = f'=IF(${AUX}{r}="","",IF(${AUX2}{r}="","（没有名称）",${AUX2}{r}))'
        ws[f'B{r}'] = f'=IF(${AUX}{r}="","",LEFT(${AUX}{r},2))'
        row(r, f'${AUX2}{r}', f'$B{r}')


# ───────────────────────── 顶部：选择格、关键数 ─────────────────────────
def top(ws):
    selector(ws, 'A3', '起（开票日期）', IN_START, None, fmt=DATE)
    selector(ws, 'A4', '止', IN_END, None, fmt=DATE)
    dv_date(ws, IN_START)
    dv_date(ws, IN_END)
    put(ws, 'C3', f'=IF(ISNUMBER({IN_START}),INT({IN_START}),P_年初)', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'C4', f'=IF(ISNUMBER({IN_END}),INT({IN_END}),P_截止)', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('D3:G3')
    ws.merge_cells('D4:G4')
    put(ws, 'D3', f'="← 实际用的起始日"&IF(ISNUMBER({IN_START}),"","（黄格空着＝首页报表年度的 1 月 1 日）")', F_NOTE, None, None, AL,
        border=False)
    put(ws, 'D4', f'=IF({S}>{E},"✗ 起比止晚，请改黄格日期","← 实际用的截止日"&IF(ISNUMBER({IN_END}),"","（黄格空着＝首页截止日期）"))',
        F_NOTE, None, None, AL, border=False)
    kp = [('H3', '销项开票\n（含税）', f'=C{R_PJ_T}'), ('H4', '销项税', f'=J{R_PJ_T}'),
          ('J3', '收到的票\n（含税）', f'=C{R_SUP_T}+C{R_SUP_C}'), ('J4', '可抵进项税\n（专票）', f'=L{R_SUP_T}'),
          ('L3', '估算应交\n增值税', f'=F{R_MON_T}'),
          ('L4', '跟税局差额', f'=IF({PIAO_N}=0,"还没粘贴",ROUND(ABS(D{R_CHK_T})+ABS(H{R_CHK_T}),2))')]
    for lc, lbl, f in kp:
        put(ws, lc, lbl, F_KPI_L, fill('FFD9E1F2'), align=ACW)
        vc = f'{CL(CI(lc[0]) + 1)}{lc[1:]}'
        put(ws, vc, f, F_KV, FILL_AUTO, MONEY, AR)
    ws.row_dimensions[3].height = ws.row_dimensions[4].height = 30
    home_link(ws, f'{LASTC}3')


TIP = ('💡 开出去的票（销项）来自【应收登记】类型「开票」；收到的票（进项）来自【应付登记】的已开票金额，再加上【收支登记】里零星采购、垫付拿到的票'
       '（已开票金额；开票日期空着按付款日期）。黄格选起止（都按开票日期算；空着＝首页的报表年度 1 月 1 日到截止日期）。'
       '① 按项目、② 按供应商分税率统计（普票、没写发票类型的票不能抵扣）；③ 逐月估算要交的增值税＝销项税−专票进项税（小规模纳税人按征收率算、不抵进项），'
       '附加税＝增值税×【基本信息】附加税率，负数＝进项多、留到下个月抵；④ 跟电子税务局导出的发票（【发票登记】）逐月对，差额不为 0 就是有票没录或记错了月份；'
       '⑤ 税局导出按对方汇总。垫付报销的已开票/欠票按人看【个人往来】。')


def build(wb, ctx=None):
    ws = wb[SH_INVSUM]
    widths(ws, {'A': 30, 'B': 12, **{CL(i): 14 for i in range(3, CI(LASTC) + 1)}, 'O': 2})
    title(ws, '发 票 统 计', LASTC, C_VIEW, TIP)
    top(ws)
    blk_out(ws)
    blk_in(ws)
    blk_month(ws)
    blk_check(ws)
    blk_cp(ws)
    for c in (IDX, AUX, AUX2, HC_PJ, HC_SUP, HC_XS, HC_JX, HC_KEY):
        ws.column_dimensions[c].hidden = True
    for col in (IDX, AUX, AUX2):
        for r in range(3, R_LAST + 1):
            if ws[f'{col}{r}'].value is not None:
                ws[f'{col}{r}'].font = F_HELP
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:1', landscape=True)
    ws.print_area = f'A1:{LASTC}{R_LAST}'
