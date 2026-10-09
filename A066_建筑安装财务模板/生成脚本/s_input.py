# -*- coding: utf-8 -*-
"""录入表（蓝）和基本信息表（灰）：只放录入列 ＋ 下拉 ＋ 少量「只给人看」的显示列（别处不引用）。
   数据来自 data_prep.build_ctx()。"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *

TODAY_HINT = '日期要像 2026/9/1 这样填'


def _d(v):
    if v in (None, ''):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.datetime(v.year, v.month, v.day)
    s = str(v)[:10]
    try:
        return dt.datetime.strptime(s, '%Y-%m-%d')
    except ValueError:
        return v


def _num(v):
    if v in (None, ''):
        return None
    try:
        return float(v) if not isinstance(v, (int, float)) else v
    except (TypeError, ValueError):
        return v


def dv_custom(ws, sqref, formula, title, msg, stop=False):
    dv = DataValidation(type='custom', formula1=formula, allow_blank=True, showErrorMessage=True,
                        errorStyle='stop' if stop else 'warning', errorTitle=title, error=msg)
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_dec(ws, sqref):
    dv = DataValidation(type='decimal', operator='between', formula1='-999999999', formula2='999999999', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle='金额', error='金额要填数字（不要带文字、单位）')
    ws.add_data_validation(dv)
    dv.add(sqref)


def _sheet_frame(ws, last_col, hdr_row, heads, color, tip, widths_, r0, n, fmts=None, aligns=None, freeze=None):
    """标题、提示、表头、数据区格式、冻结、筛选"""
    widths(ws, widths_)
    title(ws, ws.title, last_col, color, tip)
    header(ws, hdr_row, heads, color)
    style_rows(ws, r0, r0 + min(n, 300) - 1, [c for c, _ in heads], fmts=fmts or {}, aligns=aligns or {})
    ws.freeze_panes = freeze or f'B{r0}'
    ws.auto_filter.ref = f'A{hdr_row}:{last_col}{r0 + n - 1}'


# ───────────────────────── 收支登记 ─────────────────────────
def build_cash(ws, ctx):
    C = CASH_COLS
    heads = [(C['日期'], '日期'), (C['类别'], '类别'), (C['账户'], '账户'), (C['收支项目'], '收支项目'), (C['摘要'], '业务摘要'),
             (C['收入'], '收入金额'), (C['支出'], '支出金额'), (C['结余'], '账户结余\n（自动）'), (C['项目'], '项目名称'),
             (C['负责人'], '负责人'), (C['客户'], '客户名称'), (C['供应商'], '供应商名称\n（冲谁的应付）'), (C['人员'], '人员\n（工资/借款/垫付给谁）'),
             (C['对方账户'], '对方账户\n（内部转账填）'), (C['费用归属'], '费用归属\n（空＝自动）'), (C['已开票'], '已开票金额'),
             (C['开票日期'], '开票日期'), (C['备注'], '备注')]
    tip = ('💡 照你原来的「收支登记表」记：一笔一行，选类别、账户、收支项目，填摘要和金额；项目、客户、供应商、人员都从下拉选（清单在【基本信息】几张表）。'
           '支出栏可以记负数（借进来的钱、退款、冲回），跟原来一样。账户之间转钱只记一行：账户＝转出的那个、收支项目「内部转账」、对方账户＝转入的那个。'
           '老板/员工自己垫的钱：账户选他的个人账户（聂辉微信、王伟微信…）。「费用归属」空着就自动（项目上的进项目成本，办公室的进待摊），想改就选。'
           '可以插行、删行、排序（整行一起动），汇总都不受影响；只是插进来的行「账户结余」那格没公式，从上一行往下拉一下就有了。')
    W = {C['日期']: 11, C['类别']: 6, C['账户']: 13, C['收支项目']: 14, C['摘要']: 34, C['收入']: 13, C['支出']: 13, C['结余']: 14,
         C['项目']: 13, C['负责人']: 8, C['客户']: 16, C['供应商']: 16, C['人员']: 9, C['对方账户']: 12, C['费用归属']: 10,
         C['已开票']: 11, C['开票日期']: 11, C['备注']: 22}
    fm = {C['日期']: DATE, C['收入']: MONEY, C['支出']: MONEY, C['结余']: MONEY, C['已开票']: MONEY, C['开票日期']: DATE}
    al = {C['摘要']: AL, C['备注']: AL, C['收入']: AR, C['支出']: AR, C['结余']: AR, C['已开票']: AR}
    _sheet_frame(ws, CASH_LAST, CASH_HDR, heads, C_IN, tip, W, CASH_R0, N_CASH, fm, al, freeze=f'C{CASH_R0}')
    ws.row_dimensions[CASH_HDR].height = 44
    # 顶部：合计、查询（第 4～6 行）
    r1, r2 = 5, 6
    last = CASH_R0 + N_CASH - 1 + 2000
    for col, lbl, f in ((C['收入'], '收入总金额', f'=SUBTOTAL(9,{C["收入"]}{CASH_R0}:{C["收入"]}{last})'),
                        (C['支出'], '支出总金额', f'=SUBTOTAL(9,{C["支出"]}{CASH_R0}:{C["支出"]}{last})'),
                        (C['结余'], '收支盈亏', f'={C["收入"]}{r2}-{C["支出"]}{r2}')):
        put(ws, f'{col}{r1}', lbl, F_KPI_L, fill('FFD9E1F2'), align=AC)
        put(ws, f'{col}{r2}', f, F_KPI_V, FILL_AUTO, MONEY, AR)
    put(ws, f'A{r1}', '筛选后只合计看得见的行', F_NOTE, align=AL, border=False)
    ws.merge_cells(f'A{r1}:E{r1}')
    qc = [(C['客户'], '客户查询', '甲方列表', f'=SUMIFS({C["收入"]}${CASH_R0}:{C["收入"]}${last},{C["客户"]}${CASH_R0}:{C["客户"]}${last},{C["供应商"]}{r1})'
           f'-SUMIFS({C["支出"]}${CASH_R0}:{C["支出"]}${last},{C["客户"]}${CASH_R0}:{C["客户"]}${last},{C["供应商"]}{r1})', '收款合计'),
          (C['客户'], '供应商查询', '供应商列表', f'=SUMIFS({C["支出"]}${CASH_R0}:{C["支出"]}${last},{C["供应商"]}${CASH_R0}:{C["供应商"]}${last},{C["供应商"]}{r2})'
           f'-SUMIFS({C["收入"]}${CASH_R0}:{C["收入"]}${last},{C["供应商"]}${CASH_R0}:{C["供应商"]}${last},{C["供应商"]}{r2})', '付款合计')]
    for k, (lc, lbl, dvn, f, lbl2) in enumerate(qc):
        r = r1 + k
        put(ws, f'{C["客户"]}{r}', lbl, F_KPI_L, fill('FFD9E1F2'), align=AC)
        put(ws, f'{C["供应商"]}{r}', None, F_SEL, FILL_SEL, align=AC)
        dv = DataValidation(type='list', formula1=f'={dvn}', allow_blank=True, showErrorMessage=False)
        ws.add_data_validation(dv)
        dv.add(f'{C["供应商"]}{r}')
        put(ws, f'{C["人员"]}{r}', lbl2, F_KPI_L, fill('FFD9E1F2'), align=AC)
        put(ws, f'{C["对方账户"]}{r}', f'=IF({C["供应商"]}{r}="","",{f[1:]})', F_KPI_V, FILL_AUTO, MONEY, AR)
    # 数据
    rows = ctx.get('cash', [])
    assert len(rows) <= N_CASH, len(rows)
    for i, x in enumerate(rows):
        r = CASH_R0 + i
        for k in ('日期', '类别', '账户', '收支项目', '摘要', '收入', '支出', '项目', '负责人', '客户', '供应商', '人员', '对方账户',
                  '费用归属', '已开票', '开票日期', '备注'):
            v = x.get(k)
            if k in ('日期', '开票日期'):
                v = _d(v)
            elif k in ('收入', '支出', '已开票'):
                v = _num(v)
            if v not in (None, ''):
                ws[f'{C[k]}{r}'] = v
    # 账户结余（只显示；插进来的行没有也不影响任何汇总）
    hc, a, inc, out, opp = C['结余'], C['账户'], C['收入'], C['支出'], C['对方账户']
    nshow = max(len(rows) + 300, 1500)
    for i in range(min(nshow, N_CASH)):
        r = CASH_R0 + i
        ws[f'{hc}{r}'] = (f'=IF(OR(${C["日期"]}{r}="",${a}{r}=""),"",ROUND(SUMIF(账户_名称,${a}{r},账户_期初)'
                          f'+SUMIF(${a}${CASH_R0}:${a}{r},${a}{r},${inc}${CASH_R0}:${inc}{r})-SUMIF(${a}${CASH_R0}:${a}{r},${a}{r},${out}${CASH_R0}:${out}{r})'
                          f'+SUMIF(${opp}${CASH_R0}:${opp}{r},${a}{r},${out}${CASH_R0}:${out}{r})-SUMIF(${opp}${CASH_R0}:${opp}{r},${a}{r},${inc}${CASH_R0}:${inc}{r}),2))')
        ws[f'{hc}{r}'].fill = FILL_AUTO
        ws[f'{hc}{r}'].font = F_AUTO
    # 下拉
    sq = lambda col: f'{col}{CASH_R0}:{col}{CASH_R0 + N_CASH - 1}'
    dv_date(ws, sq(C['日期']))
    dv_date(ws, sq(C['开票日期']))
    dv_list(ws, sq(C['类别']), '"收入,支出"')
    dv_list(ws, sq(C['账户']), '=账户列表')
    dv_list(ws, sq(C['收支项目']), f'=INDIRECT(${C["类别"]}{CASH_R0})', prompt='先选「类别」，这里才有对应的收支项目')
    dv_list(ws, sq(C['项目']), '=项目列表')
    dv_list(ws, sq(C['客户']), '=甲方列表', stop=False)
    dv_list(ws, sq(C['供应商']), '=供应商列表', stop=False)
    dv_list(ws, sq(C['人员']), '=人员列表', stop=False)
    dv_list(ws, sq(C['对方账户']), '=账户列表')
    dv_list(ws, sq(C['费用归属']), '"' + ','.join(GUISHU) + '"')
    for col in (C['收入'], C['支出'], C['已开票']):
        dv_dec(ws, sq(col))
    # 收入行浅绿、支出行浅蓝（跟原表一样）
    rng = f'A{CASH_R0}:{CASH_LAST}{CASH_R0 + N_CASH - 1}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${C["类别"]}{CASH_R0}="收入"'], fill=fill('FFE2EFDA')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${C["类别"]}{CASH_R0}="支出"'], fill=fill('FFDDEBF7')))
    print_setup(ws, f'{CASH_HDR}:{CASH_HDR}')


# ───────────────────────── 应收登记 ─────────────────────────
def build_ar_in(ws, ctx):
    C = AR_COLS
    heads = [(C['日期'], '日期'), (C['项目'], '项目名称'), (C['客户'], '客户名称\n（空＝项目档案的甲方）'), (C['类型'], '类型'),
             (C['金额'], '金额（含税）'), (C['税率'], '税率\n（开票填）'), (C['发票号'], '发票号码'), (C['备注'], '备注')]
    tip = ('💡 每个项目一笔一行：补充协议、签证、扣款（填正数，自动减）、结算调整（多退少补，可负）、确认产值（甲方/总包批的产值）、开票（每开一张记一行：开票日期、价税合计、税率）。'
           '合同金额在【项目档案】；收款不用在这记——【收支登记】里收支项目选「工程款」、选项目就自动算进已收。【应收账款】【应收对账】自动出每个项目的应收、已收、未收、已开票、未开票。')
    W = {C['日期']: 11, C['项目']: 16, C['客户']: 24, C['类型']: 10, C['金额']: 15, C['税率']: 8, C['发票号']: 22, C['备注']: 30}
    fm = {C['日期']: DATE, C['金额']: MONEY, C['税率']: PCT}
    _sheet_frame(ws, AR_LAST, AR_HDR, heads, C_IN, tip, W, AR_R0, N_AR, fm, {C['备注']: AL, C['金额']: AR})
    for i, x in enumerate(ctx.get('ar', [])):
        r = AR_R0 + i
        for k in ('日期', '项目', '客户', '类型', '金额', '税率', '发票号', '备注'):
            v = x.get(k)
            v = _d(v) if k == '日期' else (_num(v) if k in ('金额', '税率') else v)
            if v not in (None, ''):
                ws[f'{C[k]}{r}'] = v
    sq = lambda col: f'{col}{AR_R0}:{col}{AR_R0 + N_AR - 1}'
    dv_date(ws, sq(C['日期']))
    dv_list(ws, sq(C['项目']), '=项目列表')
    dv_list(ws, sq(C['客户']), '=甲方列表', stop=False)
    dv_list(ws, sq(C['类型']), '"' + ','.join(AR_TYPES) + '"')
    dv_list(ws, sq(C['税率']), '"0.09,0.13,0.06,0.03,0.01,0"', stop=False)
    dv_dec(ws, sq(C['金额']))
    print_setup(ws, f'{AR_HDR}:{AR_HDR}')


# ───────────────────────── 应付登记 ─────────────────────────
def build_ap_in(ws, ctx):
    C = AP_COLS
    heads = [(C['日期'], '日期'), (C['项目'], '项目名称'), (C['供应商'], '供应商名称'), (C['类型'], '类型\n（空＝供应商的）'),
             (C['摘要'], '摘要（材料/台班/批次）'), (C['数量'], '数量\n（台班填小时）'), (C['单价'], '单价'), (C['应付'], '应付金额\n（空＝数量×单价）'),
             (C['已开票'], '已开票金额'), (C['开票日期'], '开票日期'), (C['税率'], '税率'), (C['发票类型'], '发票类型'), (C['备注'], '备注')]
    tip = ('💡 照你原来的应付分表记：送货、分包结算、机械台班一笔一行（日期、项目、供应商、应付金额）；收到发票就填已开票金额、开票日期、税率'
           '（可以跟应付同一行，也可以单独一行只填发票）。付款不在这记——【收支登记】付款时「供应商名称」选这家，自动冲应付。'
           '包工包料的分包：材料票 13%、人工票 1%，分两行记发票就行。【应付账款】【应付对账】自动出每家应付、已付、未付、已开票、未开票。')
    W = {C['日期']: 11, C['项目']: 14, C['供应商']: 18, C['类型']: 8, C['摘要']: 28, C['数量']: 9, C['单价']: 9, C['应付']: 14,
         C['已开票']: 13, C['开票日期']: 11, C['税率']: 7, C['发票类型']: 8, C['备注']: 22}
    fm = {C['日期']: DATE, C['数量']: '#,##0.##', C['单价']: MONEY, C['应付']: MONEY, C['已开票']: MONEY, C['开票日期']: DATE, C['税率']: PCT}
    _sheet_frame(ws, AP_LAST, AP_HDR, heads, C_IN, tip, W, AP_R0, N_AP, fm, {C['摘要']: AL, C['备注']: AL})
    for i, x in enumerate(ctx.get('ap', [])):
        r = AP_R0 + i
        for k in AP_COLS:
            v = x.get(k)
            v = _d(v) if k in ('日期', '开票日期') else (_num(v) if k in ('数量', '单价', '应付', '已开票', '税率') else v)
            if v not in (None, ''):
                ws[f'{C[k]}{r}'] = v
    sq = lambda col: f'{col}{AP_R0}:{col}{AP_R0 + N_AP - 1}'
    dv_date(ws, sq(C['日期']))
    dv_date(ws, sq(C['开票日期']))
    dv_list(ws, sq(C['项目']), '=项目列表')
    dv_list(ws, sq(C['供应商']), '=供应商列表')
    dv_list(ws, sq(C['类型']), '"' + ','.join(AP_TYPES) + '"')
    dv_list(ws, sq(C['税率']), '"0.13,0.09,0.06,0.03,0.01,0"', stop=False)
    dv_list(ws, sq(C['发票类型']), '"专票,普票"')
    for col in (C['数量'], C['单价'], C['应付'], C['已开票']):
        dv_dec(ws, sq(col))
    print_setup(ws, f'{AP_HDR}:{AP_HDR}')


# ───────────────────────── 考勤工资 ─────────────────────────
def build_att(ws, ctx):
    heads = [(ATT_MON, '月份'), (ATT_NAME, '姓名'), (ATT_TEAM, '工种/班组')]
    for i in range(NPAIR):
        heads += [(ATT_PJS[i], f'项目{i + 1}'), (ATT_DAYS[i], f'天数{i + 1}')]
    heads += [(ATT_RATE, '单价\n（手填，空＝工资标准）'), (ATT_EXTRA, '加班/补贴'), (ATT_DED, '扣款'), (ATT_NOTE, '备注'),
              (ATT_SHOW_DAYS, '总天数\n（自动）'), (ATT_SHOW_RATE, '单价\n（自动）'), (ATT_SHOW_AMT, '应发\n（自动）')]
    tip = ('💡 一人一月一行：月份（当月任一天）、姓名，然后「项目1 天数1、项目2 天数2……」成对填，一行最多 8 个工地，再多就再加一行（月薪的人按天数比例分，不会重复算）。'
           '在公司（办公室）的天数项目选「办公室」，进管理费再摊。单价空着＝【工资标准】里这个月有效的单价；临时改价在「单价」填。'
           '工资社保只是在公司账上过一下的人（比如姚俊强），在【人员信息】设成「过账人员」，就不算工资成本。右边三列只是显示，插进来的行没有也不影响汇总。')
    W = {ATT_MON: 9, ATT_NAME: 9, ATT_TEAM: 9, ATT_RATE: 11, ATT_EXTRA: 10, ATT_DED: 9, ATT_NOTE: 18, ATT_SHOW_DAYS: 8,
         ATT_SHOW_RATE: 9, ATT_SHOW_AMT: 12}
    for i in range(NPAIR):
        W[ATT_PJS[i]], W[ATT_DAYS[i]] = 12, 6
    fm = {ATT_MON: 'yyyy/m', ATT_RATE: MONEY, ATT_EXTRA: MONEY, ATT_DED: MONEY, ATT_SHOW_RATE: MONEY, ATT_SHOW_AMT: MONEY}
    for c in ATT_DAYS + [ATT_SHOW_DAYS]:
        fm[c] = '0.##'
    _sheet_frame(ws, ATT_LAST, ATT_HDR, heads, C_IN, tip, W, ATT_R0, N_ATT, fm, {ATT_NOTE: AL}, freeze=f'C{ATT_R0}')
    rows = ctx.get('att', [])
    for i, x in enumerate(rows):
        r = ATT_R0 + i
        ws[f'{ATT_MON}{r}'] = _d(x['月份'])
        ws[f'{ATT_NAME}{r}'] = x['姓名']
        if x.get('工种'):
            ws[f'{ATT_TEAM}{r}'] = x['工种']
        for j, (pj, d) in enumerate(x.get('pairs', [])[:NPAIR]):
            if pj:
                ws[f'{ATT_PJS[j]}{r}'] = pj
            if d not in (None, ''):
                ws[f'{ATT_DAYS[j]}{r}'] = _num(d)
        for k, col in (('单价', ATT_RATE), ('补贴', ATT_EXTRA), ('扣款', ATT_DED), ('备注', ATT_NOTE)):
            v = x.get(k)
            if v not in (None, '', 0):
                ws[f'{col}{r}'] = _num(v) if k != '备注' else v
    # 显示列（只给人看）
    nshow = min(max(len(rows) + 300, 800), N_ATT)
    days = ','.join(f'{c}{{r}}' for c in ATT_DAYS)
    for i in range(nshow):
        r = ATT_R0 + i
        me = f'DATE(YEAR(${ATT_MON}{r}),MONTH(${ATT_MON}{r})+1,0)'
        nm = f'${ATT_NAME}{r}'
        mode = f'IFERROR(INDEX(人员_计薪,MATCH({nm},人员_姓名,0)),"日薪")'
        std = std_rate(nm, me, f'{mode}="月薪"')
        ws[f'{ATT_SHOW_DAYS}{r}'] = f'=IF({nm}="","",SUM({days.format(r=r)}))'
        ws[f'{ATT_SHOW_RATE}{r}'] = f'=IF(OR({nm}="",NOT(ISNUMBER(${ATT_MON}{r}))),"",IF(N(${ATT_RATE}{r})<>0,${ATT_RATE}{r},IFERROR({std},0)))'
        ws[f'{ATT_SHOW_AMT}{r}'] = (f'=IF(OR({nm}="",{ATT_SHOW_RATE}{r}=""),"",ROUND(IF({mode}="月薪",{ATT_SHOW_RATE}{r}*IF(SUMIFS(${ATT_SHOW_DAYS}${ATT_R0}:${ATT_SHOW_DAYS}${ATT_R0 + nshow - 1},${ATT_NAME}${ATT_R0}:${ATT_NAME}${ATT_R0 + nshow - 1},{nm},${ATT_MON}${ATT_R0}:${ATT_MON}${ATT_R0 + nshow - 1},">="&DATE(YEAR(${ATT_MON}{r}),MONTH(${ATT_MON}{r}),1),${ATT_MON}${ATT_R0}:${ATT_MON}${ATT_R0 + nshow - 1},"<="&{me})>0,'
                                    f'N({ATT_SHOW_DAYS}{r})/SUMIFS(${ATT_SHOW_DAYS}${ATT_R0}:${ATT_SHOW_DAYS}${ATT_R0 + nshow - 1},${ATT_NAME}${ATT_R0}:${ATT_NAME}${ATT_R0 + nshow - 1},{nm},${ATT_MON}${ATT_R0}:${ATT_MON}${ATT_R0 + nshow - 1},">="&DATE(YEAR(${ATT_MON}{r}),MONTH(${ATT_MON}{r}),1),${ATT_MON}${ATT_R0}:${ATT_MON}${ATT_R0 + nshow - 1},"<="&{me}),1),'
                                    f'{ATT_SHOW_RATE}{r}*N({ATT_SHOW_DAYS}{r}))+N(${ATT_EXTRA}{r})-N(${ATT_DED}{r}),2))')
        for c in (ATT_SHOW_DAYS, ATT_SHOW_RATE, ATT_SHOW_AMT):
            ws[f'{c}{r}'].fill = FILL_AUTO
            ws[f'{c}{r}'].font = F_AUTO
    sq = lambda col: f'{col}{ATT_R0}:{col}{ATT_R0 + N_ATT - 1}'
    dv_date(ws, sq(ATT_MON))
    dv_list(ws, sq(ATT_NAME), '=人员列表')
    for c in ATT_PJS:
        dv_list(ws, sq(c), '=项目列表')
    for c in ATT_DAYS + [ATT_RATE, ATT_EXTRA, ATT_DED]:
        dv_dec(ws, sq(c))
    print_setup(ws, f'{ATT_HDR}:{ATT_HDR}')


# ───────────────────────── 发票登记（电子税务局导出粘贴区） ─────────────────────────
def build_inv(ws, ctx):
    heads = [(CL(i + 1), h) for i, h in enumerate(INV_HEAD)] + [(INV_PJ, '项目名称\n（选）')]
    tip = ('💡 电子税务局「我要办税 → 税务数字账户 → 发票查询统计 → 全量发票查询」，开票、收票各查一次点「导出」，打开导出表的「发票基础信息」页，'
           '从第一张票那行起选到最后一张（不要表头、合计），粘到 A 列第一个空行；最右边选项目。这张表只用来跟【应收登记】【应付登记】里记的票核对（【发票统计】），不重复算。')
    W = {CL(i + 1): 12 for i in range(len(INV_HEAD))}
    W.update({'F': 22, 'H': 22, 'I': 17, INV_PJ: 14})
    _sheet_frame(ws, INV_PJ, INV_HDR, heads, C_IN, tip, W, INV_R0, N_INV, {}, {})
    for r in range(INV_R0, INV_R0 + 300):
        for i in range(len(INV_HEAD)):
            ws[f'{CL(i + 1)}{r}'].fill = FILL_PASTE
    for i, row in enumerate(ctx.get('inv', [])):
        for j, v in enumerate(row):
            if v not in (None, ''):
                ws.cell(row=INV_R0 + i, column=j + 1, value=v)
    dv_list(ws, f'{INV_PJ}{INV_R0}:{INV_PJ}{INV_R0 + N_INV - 1}', '=项目列表')


# ───────────────────────── 基本信息类 ─────────────────────────
def _master(ws, ctx_rows, cols, heads, tip, W, n, fm=None, dvs=(), last=None):
    last = last or max(cols.values(), key=lambda c: CI(c))
    _sheet_frame(ws, last, M_HDR, heads, C_MASTER, tip, W, M_R0, n, fm or {}, {}, freeze=f'B{M_R0}')
    for i, x in enumerate(ctx_rows):
        r = M_R0 + i
        for k, col in cols.items():
            v = x.get(k)
            if isinstance(v, str) and len(v) >= 10 and v[4:5] == '-' and v[7:8] == '-' and k in ('开工', '完工', '质保到期', '生效'):
                v = _d(v)
            if v not in (None, ''):
                ws[f'{col}{r}'] = v
    for col, f, kw in dvs:
        dv_list(ws, f'{col}{M_R0}:{col}{M_R0 + n - 1}', f, **kw)


def build_pj(ws, ctx):
    C = PJ_COLS
    heads = [(C['名称'], '项目名称\n（简称，下拉用）'), (C['全称'], '全称 / 别名'), (C['甲方'], '甲方（客户）'), (C['合同'], '合同金额（含税）'),
             (C['保证金'], '保证金'), (C['质保比例'], '质保金比例'), (C['质保金额'], '质保金金额\n（空＝比例×应收）'), (C['质保到期'], '质保到期日'),
             (C['开工'], '开工日期'), (C['完工'], '完工日期'), (C['状态'], '状态'), (C['类型'], '类型'), (C['负责人'], '负责人'),
             (C['税率'], '税率\n（空＝9%）'), (C['备注'], '备注')]
    tip = ('💡 一个项目一行，名称用简称（各表下拉都用它）。合同金额只填原合同；补充协议、签证、扣款、结算调整在【应收登记】记。'
           '保证金是合同约定我方要交的；质保金填比例或金额、到期日。开工日期空着＝建账日期（项目账「开工至今」从开工日算起，请尽量填上）。「办公室」是公司自己（类型＝公司），不是工程项目。'
           '可以插行删行；名称别重复。')
    W = {C['名称']: 15, C['全称']: 28, C['甲方']: 24, C['合同']: 15, C['保证金']: 11, C['质保比例']: 9, C['质保金额']: 12,
         C['质保到期']: 11, C['开工']: 11, C['完工']: 11, C['状态']: 7, C['类型']: 6, C['负责人']: 8, C['税率']: 7, C['备注']: 26}
    fm = {C['合同']: MONEY, C['保证金']: MONEY, C['质保比例']: PCT, C['质保金额']: MONEY, C['质保到期']: DATE, C['开工']: DATE,
          C['完工']: DATE, C['税率']: PCT}
    _master(ws, ctx.get('projects', []), C, heads, tip, W, N_PJ, fm,
            [(C['甲方'], '=甲方列表', dict(stop=False)), (C['状态'], '"' + ','.join(PJ_STATES) + '"', {}),
             (C['类型'], '"' + ','.join(PJ_TYPES) + '"', {}), (C['税率'], '"0.09,0.13,0.06,0.03,0.01"', dict(stop=False))], PJ_LAST)


def build_cus(ws, ctx):
    C = CUS_COLS
    heads = [(C['名称'], '客户名称（甲方/总包）'), (C['简称'], '简称 / 别名'), (C['税号'], '税号'), (C['联系人'], '联系人'),
             (C['电话'], '电话'), (C['开户'], '开户行及账号 / 地址'), (C['备注'], '备注')]
    tip = '💡 甲方、总包一家一行（名称用全称，别名写在第二列）。【收支登记】客户名称、【项目档案】甲方、【应收登记】都从这里下拉。'
    _master(ws, ctx.get('customers', []), C, heads, tip, {C['名称']: 30, C['简称']: 24, C['税号']: 20, C['联系人']: 9,
                                                            C['电话']: 13, C['开户']: 30, C['备注']: 24}, N_CUS, {}, [], CUS_LAST)


def build_sup(ws, ctx):
    C = SUP_COLS
    heads = [(C['名称'], '供应商名称\n（材料商/分包/机械）'), (C['类型'], '类型'), (C['开票单位'], '开票单位全称'), (C['税号'], '税号'),
             (C['税率'], '默认税率'), (C['发票类型'], '发票类型'), (C['联系人'], '联系人'), (C['电话'], '电话'),
             (C['收款账户'], '收款账户'), (C['备注'], '备注')]
    tip = ('💡 材料商、分包、机械（运输）一家一行，类型选好（应付、成本按类型分）。名称跟你应付分表的名字一样（筑强、恩旗老叶……）；开票单位全称写在第三列。'
           '个人做分包、没有公司的，也登在这里，付款时供应商选他；给他班组工人发工资的，在【收支登记】记工资、人员选工人。')
    _master(ws, ctx.get('suppliers', []), C, heads, tip, {C['名称']: 18, C['类型']: 7, C['开票单位']: 28, C['税号']: 20, C['税率']: 8,
                                                            C['发票类型']: 8, C['联系人']: 9, C['电话']: 13, C['收款账户']: 26, C['备注']: 26},
            N_SUP, {C['税率']: PCT},
            [(C['类型'], '"' + ','.join(AP_TYPES) + '"', {}), (C['发票类型'], '"专票,普票"', {}),
             (C['税率'], '"0.13,0.09,0.06,0.03,0.01,0"', dict(stop=False))], SUP_LAST)


def build_per(ws, ctx):
    C = PER_COLS
    heads = [(C['姓名'], '姓名'), (C['类别'], '类别'), (C['工种'], '工种/班组'), (C['计薪'], '计薪方式'), (C['过账'], '过账人员\n（是＝不算工资成本）'),
             (C['期初欠薪'], '建账前欠薪'), (C['身份证'], '身份证号'), (C['电话'], '电话'), (C['银行卡'], '银行卡'), (C['备注'], '备注 / 别名')]
    tip = ('💡 工人、班组长、管理人员、老板一人一行。计薪方式：日薪（按天）或月薪。「过账人员」＝工资、社保只是在公司账上过一下、以后要退回的人（比如姚俊强）：'
           '选「是」以后，他的考勤不算工资成本，公司替他付的、他退回来的都记在【个人往来】「代收代付」里。')
    _master(ws, ctx.get('persons', []), C, heads, tip, {C['姓名']: 10, C['类别']: 9, C['工种']: 10, C['计薪']: 8, C['过账']: 10,
                                                          C['期初欠薪']: 11, C['身份证']: 20, C['电话']: 13, C['银行卡']: 22, C['备注']: 26},
            N_PER, {C['期初欠薪']: MONEY},
            [(C['类别'], '"' + ','.join(PER_KINDS) + '"', {}), (C['计薪'], '"' + ','.join(PAY_MODES) + '"', {}),
             (C['过账'], '"是,否"', {})], PER_LAST)


def build_rate(ws, ctx):
    C = RATE_COLS
    heads = [(C['姓名'], '姓名'), (C['生效'], '从哪天起'), (C['日单价'], '日单价（日薪）'), (C['月薪'], '月薪（月薪的人）'), (C['备注'], '备注')]
    tip = '💡 涨工资不改原来那行，加一行「从哪天起、新单价」。考勤按那个月最后一天有效的单价算，以前的月份不变。'
    _master(ws, ctx.get('rates', []), C, heads, tip, {C['姓名']: 10, C['生效']: 12, C['日单价']: 13, C['月薪']: 13, C['备注']: 26},
            N_RATE, {C['生效']: DATE, C['日单价']: MONEY, C['月薪']: MONEY}, [(C['姓名'], '=人员列表', dict(stop=False))], RATE_LAST)
    dv_date(ws, f"{C['生效']}{M_R0}:{C['生效']}{M_R0 + N_RATE - 1}")


def build_open(ws, ctx):
    C = OPEN_COLS
    heads = [(C['类型'], '类型'), (C['对象'], '对象（项目/供应商/人员/账户）'), (C['项目'], '项目'), (C['金额'], '金额'), (C['说明'], '说明')]
    tip = ('💡 建账日（【基本信息】①）以前就有的余额，一笔一行。现在建账日是 2023-05-01，所有流水都从那以后记，所以这里是空的。'
           '金额正负：应收、应付、欠薪、借款、垫付都填正数（欠我们的/我们欠的都是正）；未分配利润填以前年度累计利润（亏损填负数）。')
    _master(ws, ctx.get('opening', []), C, heads, tip, {C['类型']: 12, C['对象']: 22, C['项目']: 14, C['金额']: 15, C['说明']: 40},
            N_OPEN, {C['金额']: MONEY}, [(C['类型'], '"' + ','.join(OPEN_TYPES) + '"', {}), (C['项目'], '=项目列表', {})], OPEN_LAST)


def build_base(ws, ctx):
    widths(ws, {'A': 16, 'B': 22, 'C': 26, 'D': 14, 'E': 22, 'F': 3, 'G': 16, 'H': 13, 'I': 3, 'J': 18, 'K': 13, 'L': 11})
    tip = ('💡 ① 参数；② 账户（银行、现金、专户、个人账户：个人账户是老板/员工自己的微信、银行卡，垫付、代收都记在这里，选「所属人」）；'
           '③④ 收支项目（收支登记的下拉）：可以加、可以改名，但每个都要选一个「归类」——报表按归类汇总（归类的意思看 ⑤）。')
    title(ws, ws.title, BASE_LAST, C_MASTER, tip)
    section(ws, BASE_PAR_R0 - 1, 'A', 'C', '① 参数', C_MASTER)
    pv = ctx.get('params', {})
    for key, lbl, default, fmt, dvs in PARAMS:
        r = PAR_ROW[key]
        put(ws, f'B{r}', lbl, F_TXTB, fill('FFD9E1F2'), align=AL)
        v = pv.get(key, default)
        if fmt == 'date':
            v = _d(v)
        put(ws, f'C{r}', v, F_IN, FILL_IN, fmt={'date': DATE, 'pct': PCT}.get(fmt), align=AC)
        if dvs:
            dv_list(ws, f'C{r}', f'"{dvs}"')
    section(ws, BASE_ACC_HDR - 1, 'A', 'E', '② 账户', C_MASTER)
    header(ws, BASE_ACC_HDR, [('A', '账户名称'), ('B', '类型'), ('C', '所属人（个人账户）'), ('D', '建账日余额'), ('E', '备注')], C_MASTER)
    section(ws, BASE_INC_HDR - 1, 'G', 'H', '③ 收入类收支项目', C_MASTER)
    header(ws, BASE_INC_HDR, [('G', '收支项目'), ('H', '归类')], C_MASTER)
    section(ws, BASE_EXP_HDR - 1, 'J', 'L', '④ 支出类收支项目', C_MASTER)
    header(ws, BASE_EXP_HDR, [('J', '收支项目'), ('K', '归类'), ('L', '默认费用归属\n（空＝按归类）')], C_MASTER)
    style_rows(ws, BASE_ACC_R0, BASE_ACC_R0 + N_ACC - 1, list('ABCDE'), fmts={'D': MONEY})
    style_rows(ws, BASE_INC_R0, BASE_INC_R0 + N_INC - 1, list('GH'))
    style_rows(ws, BASE_EXP_R0, BASE_EXP_R0 + N_EXP - 1, list('JKL'))
    for i, x in enumerate(ctx.get('accounts', [])):
        r = BASE_ACC_R0 + i
        for k, col in ACC_COLS.items():
            v = x.get(k)
            if v not in (None, ''):
                ws[f'{col}{r}'] = v
    for i, x in enumerate(ctx.get('inc_items', [])):
        ws[f'G{BASE_INC_R0 + i}'], ws[f'H{BASE_INC_R0 + i}'] = x['名称'], x['归类']
    for i, x in enumerate(ctx.get('exp_items', [])):
        r = BASE_EXP_R0 + i
        ws[f'J{r}'], ws[f'K{r}'] = x['名称'], x['归类']
        if x.get('归属'):
            ws[f'L{r}'] = x['归属']
    cat_rng = f'$A${BASE_CAT_R0}:$A${BASE_CAT_R0 + len(CATS) - 1}'
    dv_list(ws, f'B{BASE_ACC_R0}:B{BASE_ACC_R0 + N_ACC - 1}', '"' + ','.join(ACC_TYPES) + '"')
    dv_list(ws, f'C{BASE_ACC_R0}:C{BASE_ACC_R0 + N_ACC - 1}', '=人员列表', stop=False)
    dv_list(ws, f'H{BASE_INC_R0}:H{BASE_INC_R0 + N_INC - 1}', f'={cat_rng}')
    dv_list(ws, f'K{BASE_EXP_R0}:K{BASE_EXP_R0 + N_EXP - 1}', f'={cat_rng}')
    dv_list(ws, f'L{BASE_EXP_R0}:L{BASE_EXP_R0 + N_EXP - 1}', '"' + ','.join(GUISHU) + '"')
    # ⑤ 归类说明（系统固定，别改）
    section(ws, BASE_CAT_HDR - 1, 'A', 'E', '⑤ 归类（系统固定，别改；③④ 的「归类」从这里选）', C_MASTER)
    header(ws, BASE_CAT_HDR, [('A', '归类'), ('B', '方向'), ('C', '进报表哪里'), ('D', '默认费用归属'), ('E', '')], C_MASTER)
    for i, (nm, dr, desc, gs, _c) in enumerate(CATS):
        r = BASE_CAT_R0 + i
        put(ws, f'A{r}', nm, F_TXTB, FILL_AUTO, align=AL)
        put(ws, f'B{r}', dr, F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'C{r}', desc, F_AUTO, FILL_AUTO, align=ALW)
        ws.merge_cells(f'C{r}:E{r}')
        put(ws, f'D{r}', None, F_AUTO, FILL_AUTO)
        ws.row_dimensions[r].height = 30
    # 归类默认归属放在隐藏列 M（_表 取用）
    for i, (nm, dr, desc, gs, _c) in enumerate(CATS):
        ws[f'M{BASE_CAT_R0 + i}'] = nm
        ws[f'N{BASE_CAT_R0 + i}'] = gs
    hide(ws, 'M', 'N')
    ws.freeze_panes = 'A4'


BUILDERS = {SH_CASH: build_cash, SH_AR_IN: build_ar_in, SH_AP_IN: build_ap_in, SH_ATT: build_att, SH_INV: build_inv,
            SH_PJ: build_pj, SH_CUS: build_cus, SH_SUP: build_sup, SH_PER: build_per, SH_RATE: build_rate, SH_BASE: build_base,
            SH_OPEN: build_open}


def build(wb, ctx):
    for name, fn in BUILDERS.items():
        fn(wb[name], ctx)
