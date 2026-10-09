# -*- coding: utf-8 -*-
"""报表组：【利润表】【资产负债表】【盈亏平衡表】（口径见 报表口径.md §3、§5、§6、§8、§9、§10）。

   三张表用同一套「某段时间 (a, b] 的利润表各行」公式（lines），资产负债表的「建账以来累计利润」也是它（a 不设下限），
   所以 利润表本年累计 ＝ 资产负债表 未分配利润(期末) − 未分配利润(年初)（期末＝截止日时），两边一定对得上。

   每一笔成本费用只进一行：
     应付登记 → 按 付_类型 进 材料/分包/机械/其他直接费（全部，含没项目的）；
     收支登记 → 收_成本类 × 收_归属：归属＝项目成本 的进营业成本（成本类「管理费用」的进其他直接费），
                归属＝待摊费用 的全部、归属＝公司费用 且成本类不是 财务费用/营业外支出/未分类 的 → 管理费用；
                财务费用、营业外支出、未分类 各自一行（这三类 _收 一律是公司费用）；
     考勤 → 计提的应发：分到工程项目的进人工（＝应发 − 未分摊），未分摊（办公室、管理人员）进管理费用。
   冲应付、冲工资、过账的收支行 收_成本类 为空，不进利润表（进资产负债表的应付、应付工资、代收代付）。"""
from layout import *
from common import *
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill

COSTS = ['材料', '分包', '机械', '人工', '其他直接费']                      # 营业成本各行
CASH8 = ['材料', '分包', '机械', '人工', '其他直接费', '管理费用', '财务费用', '营业外支出']   # 收_成本类（未分类单列）
F_RED_NOTE = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_OK = Font(name=YH, sz=10, bold=True, color='FF375623')
CF_WARN = PatternFill('solid', start_color='FFFFC7CE', end_color='FFFFC7CE')   # 条件格式的底色（Excel 认 bgColor，两个都给）
CF_OK = PatternFill('solid', start_color='FFE2EFDA', end_color='FFE2EFDA')
C_LBL = fill('FFD9E1F2')


# ═══════════════════════════ 公式片段 ═══════════════════════════
def _crit(col, a, b):
    """日期段 (a, b]；a=None＝不设下限（建账以来累计，跟 _汇「<=d」一样连没填日期的有效行也算上）"""
    s = f'{col},"<="&{b}'
    return f'{col},">"&{a},{s}' if a is not None else s


def lines(a, b):
    """(a, b] 期间全公司利润表各行（不含工程收入，它要从 _汇 取）。都是正数＝费用（未分类、其他收入正数＝收入）。"""
    S, F, FK = _crit('收_日期', a, b), _crit('付_日期', a, b), _crit('付_开票日期', a, b)
    Y, K = _crit('应_日期', a, b), _crit('考_月', a, b)

    def cash(cat, gs=None):
        g = f'收_归属,"{gs}",' if gs else ''
        return f'SUMIFS(收_净额,收_成本类,"{cat}",{g}{S})'

    def ap(t):
        return f'SUMIFS(付_应付额,付_类型,"{t}",{F})'

    kq_all = f'SUMIFS(考_应发,考_计提,1,{K})'
    kq_co = f'SUMIFS(考_未分摊,考_计提,1,{K})'
    d = {
        '其他收入': f'SUMIFS(收_净额,收_归类,"其他收入",{S})',
        '材料': f'{ap("材料")}-{cash("材料", "项目成本")}',
        '分包': f'{ap("分包")}-{cash("分包", "项目成本")}',
        '机械': f'{ap("机械")}-{cash("机械", "项目成本")}',
        '人工': f'{kq_all}-{kq_co}-{cash("人工", "项目成本")}',
        '其他直接费': f'{ap("其他")}-{cash("其他直接费", "项目成本")}-{cash("管理费用", "项目成本")}',
        '税金': f'(SUMIFS(应_销项税,{Y})-SUMIFS(付_进项税,{FK}))*(1+P_附加税率)',
        '管理费用': (f'{kq_co}-SUMIFS(收_净额,收_归属,"待摊费用",{S})-SUMIFS(收_净额,收_归属,"公司费用",{S})'
                 f'+{cash("财务费用")}+{cash("营业外支出")}+{cash("未分类")}'),
        '财务费用': f'-{cash("财务费用")}',
        '营业外支出': f'-{cash("营业外支出")}',
        '未分类': cash('未分类'),
        # 核对：所有成本费用逐笔加总（应付登记全部 ＋ 考勤计提应发全部 ＋ 收支登记 8 个成本类全部）
        '总额': f'SUMIFS(付_应付额,{F})+{kq_all}-(' + '+'.join(cash(x) for x in CASH8) + ')',
    }
    return d


def direct_cost(L):
    return '+'.join(f'({L[k]})' for k in COSTS)


def op(t):
    """期初余额表某类型合计（正数＝该科目正常方向）"""
    return f'SUMIFS(期初_金额,期初_类型,"{t}")'


def _cell(ws, coord, v, font=F_AUTO, fl=None, fmt=MONEY, al=AR):
    put(ws, coord, v, font, fl, fmt, al)


def _note(ws, r, c1, c2, text, font=F_NOTE, h=None):
    ws.merge_cells(f'{c1}{r}:{c2}{r}')
    put(ws, f'{c1}{r}', text, font, align=ALW, border=False)
    if h:
        ws.row_dimensions[r].height = h


def _sel_row(ws, start_lbl='本期 起', default_txt='空着＝首页报表年度的年初～截止日'):
    """第 3 行：起、止黄格（C3、E3）；第 4 行：实际用的日期"""
    selector(ws, 'B3', start_lbl, 'C3', None, fmt=DATE)
    selector(ws, 'D3', '止', 'E3', None, fmt=DATE)
    dv_date(ws, 'C3')
    dv_date(ws, 'E3')
    put(ws, 'B4', '实际用', F_NOTE, align=AC)
    put(ws, 'D4', '～', F_NOTE, align=AC)
    ws.merge_cells('F3:I3')
    put(ws, 'F3', f'← 填日期（{default_txt}）', F_NOTE, align=AL, border=False)


# ═══════════════════════════ 利润表 ═══════════════════════════
PL_R = {}


def build_pl(ws):
    MC = [CL(4 + i) for i in range(12)]          # D..O ＝ 1～12 月
    C_ALL = 'P'
    last = C_ALL
    widths(ws, {'A': 40, 'B': 14, 'C': 14, **{c: 12.5 for c in MC}, C_ALL: 15, 'Q': 9})
    tip = ('💡 全公司赚了多少钱。「本期」在 C3、E3 两个黄格填起止日期（空着＝首页报表年度的年初～截止日）；右边是本年累计、报表年度每个月、建账以来累计。'
           '收入按【基本信息】的「收入确认口径」算（默认：每个项目开票和收款哪个大算哪个），成本按发生算：应付登记挂账的、考勤应发的工资都算进成本，'
           '付款、发工资时就不再算一遍。每一笔成本费用只进一行（最下面有核对）。金额都含税，税金是按发票估算的。')
    title(ws, '利 润 表（全公司 · 管理用）', last, C_RPT, tip)
    home_link(ws, 'Q1')
    _sel_row(ws)
    put(ws, 'C4', '=汇_利润起', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'E4', '=汇_利润止', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('K3:P3')
    put(ws, 'K3', '="收入确认口径："&P_收入口径&"　｜　"&P_纳税人&"　｜　附加税率 "&TEXT(P_附加税率,"0%")', F_KPI_L, align=AL, border=False)
    ws.merge_cells('F4:P4')
    put(ws, 'F4', '="每个月＝"&P_年度&" 年各月（超过截止日 "&TEXT(P_截止,"yyyy-mm-dd")&" 的月份是 0）；建账以来＝建账日～"&TEXT(汇_日12月,"yyyy-mm-dd")',
        F_NOTE, align=AL, border=False)

    H = 5
    hdr = [('A', '项    目'), ('B', '本期'), ('C', '=P_年度&"年累计"')]
    hdr += [(c, f'=P_年度&"年{i + 1}月"') for i, c in enumerate(MC)]
    hdr += [(C_ALL, '建账以来累计')]
    header(ws, H, hdr, C_RPT, height=30)

    # 各列的期间 (a, b] 和工程收入
    cols = [('B', '汇_日利润起前', '汇_日利润止', '汇_收入利润止-汇_收入利润起前'),
            ('C', '汇_日年初前', '汇_日12月', '汇_收入12月-汇_收入年初前')]
    for i, c in enumerate(MC):
        pa = '年初前' if i == 0 else f'{i}月'
        cols.append((c, f'汇_日{pa}', f'汇_日{i + 1}月', f'汇_收入{i + 1}月-汇_收入{pa}'))
    cols.append((C_ALL, None, '汇_日12月', '汇_收入12月'))

    rows = [  # (键, 标签, 种类)
        ('收入', '一、收入合计', 'sec'),
        ('工程收入', '　　工程收入（按收入确认口径）', 'line'),
        ('其他收入', '　　其他收入（利息收入、杂项）', 'line'),
        ('营业成本', '二、营业成本（直接成本）', 'sec'),
        ('材料', '　　材料', 'line'),
        ('分包', '　　分包', 'line'),
        ('机械', '　　机械（台班、运输）', 'line'),
        ('人工', '　　人工（工地工人工资）', 'line'),
        ('其他直接费', '　　其他直接费（项目上的零星费用）', 'line'),
        ('毛利', '毛利（一 − 二）', 'tot'),
        ('毛利率', '　　毛利率（毛利 ÷ 收入）', 'pct'),
        ('税金', '三、税金及附加（估算）', 'sec1'),
        ('期间费用', '四、期间费用', 'sec'),
        ('管理费用', '　　管理费用（办公、招待、车辆、管理人员工资…）', 'line'),
        ('财务费用', '　　财务费用（利息、手续费）', 'line'),
        ('营业外支出', '　　营业外支出（罚款、滞纳金）', 'line'),
        ('未分类', '五、未分类收支（收入为正、支出为负，去收支登记改归类）', 'sec1'),
        ('净利润', '六、净利润（毛利 − 税金 − 期间费用 ＋ 未分类）', 'grand'),
        ('净利率', '　　净利率（净利润 ÷ 收入）', 'pct'),
        (None, None, None),
        ('总额', '核对：所有成本费用逐笔加总（应付登记＋考勤应发＋收支登记成本费用）', 'chk'),
        ('差额', '核对：逐笔加总 − 上面二、四两项各行之和（应为 0）', 'chk'),
    ]
    r = H + 1
    for k, _l, _t in rows:
        if k:
            PL_R[k] = r
        r += 1
    R = PL_R
    for k, lab, kind in rows:
        if not k:
            continue
        r = R[k]
        bold = kind in ('sec', 'sec1', 'tot', 'grand')
        fl = {'sec': FILL_SUB, 'tot': FILL_TOT, 'grand': FILL_TOT}.get(kind)
        put(ws, f'A{r}', lab, F_TXTB if bold else (F_NOTE if kind == 'chk' else F_TXT), fl, align=AL)
        for c, a, b, rev in cols:
            L = lines(a, b)
            if k == '工程收入':
                f = rev
            elif k in L:
                f = L[k]
            elif k == '收入':
                f = f'{c}{R["工程收入"]}+{c}{R["其他收入"]}'
            elif k == '营业成本':
                f = f'SUM({c}{R["材料"]}:{c}{R["其他直接费"]})'
            elif k == '毛利':
                f = f'{c}{R["收入"]}-{c}{R["营业成本"]}'
            elif k == '毛利率':
                f = f'IF(N({c}{R["收入"]})=0,"",{c}{R["毛利"]}/{c}{R["收入"]})'
            elif k == '期间费用':
                f = f'SUM({c}{R["管理费用"]}:{c}{R["营业外支出"]})'
            elif k == '净利润':
                f = f'{c}{R["毛利"]}-{c}{R["税金"]}-{c}{R["期间费用"]}+{c}{R["未分类"]}'
            elif k == '净利率':
                f = f'IF(N({c}{R["收入"]})=0,"",{c}{R["净利润"]}/{c}{R["收入"]})'
            elif k == '差额':
                f = f'ROUND({c}{R["总额"]}-{c}{R["营业成本"]}-{c}{R["期间费用"]},2)'
            font = F_AUTOB if bold else (F_NOTE if kind == 'chk' else F_AUTO)
            _cell(ws, f'{c}{r}', f'={f}', font, fl, PCT if kind == 'pct' else MONEY)
    rd = R['差额']
    ws.conditional_formatting.add(f'B{rd}:{C_ALL}{rd}', FormulaRule(formula=[f'ROUND(B{rd},2)<>0'], font=F_RED_NOTE, fill=CF_WARN))
    ws.conditional_formatting.add(f'B{R["未分类"]}:{C_ALL}{R["未分类"]}',
                                  FormulaRule(formula=[f'ROUND(B{R["未分类"]},2)<>0'], fill=CF_WARN))
    r = rd + 1
    ws.merge_cells(f'A{r}:{C_ALL}{r}')
    put(ws, f'A{r}', (f'=IF(AND(ROUND(B{rd},2)=0,ROUND(C{rd},2)=0,ROUND({C_ALL}{rd},2)=0),'
                      f'"✓ 每一笔成本费用都进了利润表，而且只进了一行",'
                      f'"✗ 有成本费用没进任何一行（差额见上一行）：多半是【应付登记】的类型不是 材料/分包/机械/其他，或【收支登记】的费用归属不对，去【数据校验】看")'
                      f'&IF(ROUND(B{R["未分类"]},2)<>0,"；⚠ 本期有未分类收支 "&TEXT(B{R["未分类"]},"#,##0.00")&" 元：【收支登记】最右边校验列有 ✗ 的行，选对收支项目（或内部转账填对方账户）就没了","")'),
        F_TXTB, align=ALW, border=False)
    ws.row_dimensions[r].height = 30
    PL_R['状态'] = r

    # 说明
    r += 2
    _note(ws, r, 'A', C_ALL, '说明', F_TXTB)
    notes = [
        ('收入口径（【基本信息】①「收入确认口径」，▶ 是现在用的）：', None),
        ('开票与收款取大', '开票与收款取大：每个项目累计开了多少票、累计收了多少工程款，哪个大算哪个；甲方批了产值（【应收登记】「确认产值」）的项目按批的产值算。'),
        ('开票', '开票：开了多少票就算多少收入（【应收登记】类型「开票」）。'),
        ('收款', '收款：收到多少工程款就算多少收入（【收支登记】收支项目归类「工程款收款」，专户、个人账户收的也算）。'),
        ('确认产值', '确认产值：只按甲方/总包批的产值算（【应收登记】「确认产值」）；没批产值的项目收入是 0。'),
        (None, '收入都是按项目累计算的：某段时间的收入＝到期末的累计收入 − 到期初前一天的累计收入，所以一个月一个月加起来正好等于全年。'
               '其他收入＝收支项目归类「其他收入」的（利息收入、杂项）。'),
        (None, '金额都含税（跟收支登记一样）。税金及附加是估算的：销项税（开票额 × 税率 ÷ (1＋税率)）− 进项税（应付登记里的专票），再乘 (1＋附加税率)；'
               '小规模纳税人按征收率算、不抵进项。实际交多少以报税为准，交税的钱记在收支登记「税费」，冲应交税费，不再进利润表。'),
        (None, '成本按发生算：【应付登记】记的送货、分包结算、台班按类型进材料/分包/机械/其他直接费（挂账没付的也算，付款时冲应付不再算）；'
               '没登记应付、直接付钱的（零星材料、个人班组）按收支项目的归类进对应的行；考勤工资按考勤应发算（考勤起算月以后），'
               '工地上的进人工，办公室/管理人员的进管理费用；发工资时冲应付工资，不再算一遍。'),
        (None, '管理费用＝费用归属「待摊费用」的全部 ＋「公司费用」里不是财务费用/营业外支出的 ＋ 考勤里没分到工程项目的工资。'
               '借款、还款、保证金、备用金、报销还款、内部转账、股东投入、买固定资产都不是成本，不进这张表（在【资产负债表】）。'),
    ]
    for key, text in notes:
        r += 1
        if key is None and text is None:
            continue
        if text is None:
            _note(ws, r, 'A', C_ALL, key, F_TXTB)
        elif key:
            _note(ws, r, 'A', C_ALL, f'=IF(P_收入口径="{key}","▶ ","　 ")&"{text}"', F_NOTE)
        else:
            _note(ws, r, 'A', C_ALL, text, F_NOTE, h=30)
    ws.freeze_panes = f'B{H + 1}'
    print_setup(ws, f'{H}:{H}', landscape=True)


# ═══════════════════════════ 资产负债表 ═══════════════════════════
BS_R = {}


def build_bs(ws):
    last = 'F'
    widths(ws, {'A': 34, 'B': 16, 'C': 16, 'D': 36, 'E': 16, 'F': 16, 'G': 9, 'H': 3, 'I': 30, 'J': 15, 'K': 15})
    tip = ('💡 公司现在有多少家底。左边是公司有的（钱、别人欠的），右边是欠别人的和股东自己的（投进来的本钱＋攒下的利润）。'
           'C3 黄格填日期（空着＝首页截止日），「年初」列是报表年度上一年的年末。两边必须相等：最下面「核对」两格都是 0 才对，不是 0 说明有一类钱没归进来。')
    title(ws, '资 产 负 债 表（全公司 · 管理用）', last, C_RPT, tip)
    home_link(ws, 'G1')
    selector(ws, 'B3', '资产负债表日', 'C3', None, fmt=DATE)
    dv_date(ws, 'C3')
    ws.merge_cells('D3:F3')
    put(ws, 'D3', '← 填日期（空着＝首页截止日）', F_NOTE, align=AL, border=False)
    put(ws, 'B4', '实际用', F_NOTE, align=AC)
    put(ws, 'C4', '=汇_负债日', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('D4:F4')
    put(ws, 'D4', '="年初列＝"&TEXT(汇_日年初前,"yyyy-mm-dd")&"（"&P_年度&" 年年初的前一天）"', F_NOTE, align=AL, border=False)
    H = 5
    header(ws, H, [('A', '资    产'), ('B', '="年初"&"（"&TEXT(汇_日年初前,"yyyy/m/d")&"）"'),
                   ('C', '="期末"&"（"&TEXT(汇_负债日,"yyyy/m/d")&"）"'),
                   ('D', '负债和所有者权益'), ('E', '="年初"&"（"&TEXT(汇_日年初前,"yyyy/m/d")&"）"'),
                   ('F', '="期末"&"（"&TEXT(汇_负债日,"yyyy/m/d")&"）"')], C_RPT, height=34)

    # ── 隐藏计算列 I:K（I 名称、J 年初前、K 负债日）：建账以来累计利润（跟利润表同一套公式）＋ 几个净额 ──
    DT = {'J': ('汇_日年初前', '年初前'), 'K': ('汇_日负债日', '负债日')}
    HK = ['工程收入', '其他收入'] + COSTS + ['税金', '管理费用', '财务费用', '营业外支出', '未分类', '累计净利润',
                                       '期初未分配', '往来净额', '应交税费净额', '期初表未分配', '核对总额', '核对差额']
    HR = {k: H + 1 + i for i, k in enumerate(HK)}
    put(ws, f'I{H}', '（隐藏）建账以来累计到该日', F_NOTE, border=False)
    put(ws, f'J{H}', '年初前', F_NOTE, border=False)
    put(ws, f'K{H}', '负债日', F_NOTE, border=False)
    init_rp = ('SUM(账户_期初)+' + '+'.join(op(t) for t in ('应收账款', '备用金', '保证金押金', '代收代付', '其他应收', '固定资产'))
               + '-(' + '+'.join(op(t) for t in ('应付账款', '应付工资', '个人借款', '个人垫付', '其他应付', '短期借款', '应交税费'))
               + ')-SUM(人员_期初欠薪)-' + op('实收资本'))
    for c, (d, suf) in DT.items():
        L = lines(None, d)
        for k in HK:
            r = HR[k]
            if k == '工程收入':
                f = f'汇_收入{suf}'
            elif k in L:
                f = L[k]
            elif k == '累计净利润':
                f = (f'{c}{HR["工程收入"]}+{c}{HR["其他收入"]}-SUM({c}{HR["材料"]}:{c}{HR["其他直接费"]})-{c}{HR["税金"]}'
                     f'-{c}{HR["管理费用"]}-{c}{HR["财务费用"]}-{c}{HR["营业外支出"]}+{c}{HR["未分类"]}')
            elif k == '期初未分配':
                f = init_rp
            elif k == '往来净额':
                f = f'{op("其他应收")}-{op("其他应付")}-SUMIFS(收_净额,收_归类,"往来款",收_日期,"<="&{d})'
            elif k == '应交税费净额':
                f = f'{op("应交税费")}+{c}{HR["税金"]}+SUMIFS(收_净额,收_归类,"税费",收_日期,"<="&{d})'
            elif k == '期初表未分配':
                f = op('未分配利润')
            elif k == '核对总额':
                f = L['总额']
            elif k == '核对差额':
                f = (f'ROUND({c}{HR["核对总额"]}-SUM({c}{HR["材料"]}:{c}{HR["其他直接费"]})-{c}{HR["管理费用"]}'
                     f'-{c}{HR["财务费用"]}-{c}{HR["营业外支出"]},2)')
            ws[f'{c}{r}'] = f'={f}'
            ws[f'{c}{r}'].number_format = MONEY
            ws[f'{c}{r}'].font = F_NOTE
            ws[f'I{r}'] = k
            ws[f'I{r}'].font = F_NOTE
    hide(ws, 'I', 'J', 'K')

    # ── 表体 ──
    def cum(cat, d):
        return f'SUMIFS(收_净额,收_归类,"{cat}",收_日期,"<="&{d})'

    def hv(k, hc):
        return f'{hc}{HR[k]}'

    left = [  # (行, 键, 标签, 种类, 公式(d, suf, hc))
        ('货币资金', '货币资金（银行＋现金＋专户）', 'line', lambda d, s, hc: f'汇_货币资金{s}'),
        ('应收账款', '应收账款（甲方欠的工程款，按收入口径）', 'line', lambda d, s, hc: f'汇_应收{s}'),
        ('预付账款', '预付账款（多付给供应商的）', 'line', lambda d, s, hc: f'汇_预付{s}'),
        ('其他应收', '其他应收款', 'sec', None),
        ('个人持有', '　　个人手上的公司钱（个人账户结余）', 'line', lambda d, s, hc: f'汇_个人持有{s}'),
        ('保证金', '　　保证金、押金', 'line', lambda d, s, hc: f'{op("保证金押金")}-{cum("保证金押金", d)}'),
        ('备用金', '　　备用金', 'line', lambda d, s, hc: f'{op("备用金")}-{cum("备用金", d)}'),
        ('代收代付', '　　代收代付（替人垫的社保、过账工资）', 'line',
         lambda d, s, hc: f'{op("代收代付")}-{cum("代收代付", d)}-SUMIFS(收_净额,收_归类,"工资发放",收_过账,1,收_日期,"<="&{d})'),
        ('往来应收', '　　往来款（别人欠我们的，净额）', 'line', lambda d, s, hc: f'MAX(0,{hv("往来净额", hc)})'),
        ('留抵', '应交税费留抵（税交多了 / 进项多）', 'line', lambda d, s, hc: f'MAX(0,-{hv("应交税费净额", hc)})'),
        ('固定资产', '固定资产（买价，不提折旧）', 'line', lambda d, s, hc: f'{op("固定资产")}-{cum("固定资产购置", d)}'),
    ]
    right = [
        ('应付账款', '应付账款（欠供应商的）', 'line', lambda d, s, hc: f'汇_应付{s}'),
        ('预收账款', '预收账款（工程款多收的）', 'line', lambda d, s, hc: f'汇_预收{s}'),
        ('应付工资', '应付工资（欠工人的）', 'line',
         lambda d, s, hc: (f'SUM(人员_期初欠薪)+{op("应付工资")}+SUMIFS(考_应发,考_计提,1,考_月,"<="&{d})'
                           f'+SUMIFS(收_净额,收_冲工资,1,收_日期,"<="&{d})')),
        ('其他应付', '其他应付款', 'sec', None),
        ('个人借款', '　　个人借款（借老板、个人的钱）', 'line', lambda d, s, hc: f'{op("个人借款")}+{cum("个人借款", d)}'),
        ('个人垫付', '　　个人垫付还没报销的', 'line', lambda d, s, hc: f'汇_个人垫付{s}'),
        ('往来应付', '　　往来款（我们欠别人的，净额）', 'line', lambda d, s, hc: f'MAX(0,-{hv("往来净额", hc)})'),
        ('短期借款', '短期借款（银行贷款）', 'line', lambda d, s, hc: f'{op("短期借款")}+{cum("银行借款", d)}'),
        ('应交税费', '应交税费（估算还没交的）', 'line', lambda d, s, hc: f'MAX(0,{hv("应交税费净额", hc)})'),
        ('负债合计', '负债合计', 'tot', None),
        ('实收资本', '实收资本（股东投进来的本钱）', 'line', lambda d, s, hc: f'{op("实收资本")}+{cum("股东投入", d)}'),
        ('未分配', '未分配利润（攒下来的利润）', 'sec', None),
        ('期初未分配', '　　其中：建账前的（按期初余额自动倒推）', 'sub', lambda d, s, hc: hv('期初未分配', hc)),
        ('累计利润', '　　　　　建账以来累计利润（同利润表算法）', 'sub', lambda d, s, hc: hv('累计净利润', hc)),
        ('权益合计', '所有者权益合计', 'tot', None),
    ]
    r0 = H + 1
    nrow = max(len(left), len(right))
    RT = r0 + nrow                                   # 合计行
    for i, x in enumerate(left):
        BS_R['L_' + x[0]] = r0 + i
    for i, x in enumerate(right):
        BS_R['R_' + x[0]] = r0 + i
    BS_R['资产合计'] = BS_R['负债和权益合计'] = RT
    R = BS_R
    VC = {'L': [('B', 'J'), ('C', 'K')], 'R': [('E', 'J'), ('F', 'K')]}
    for side, items, lc in (('L', left, 'A'), ('R', right, 'D')):
        for i in range(nrow):
            r = r0 + i
            if i >= len(items):
                for c in [lc] + [v for v, _h in VC[side]]:
                    put(ws, f'{c}{r}', None)
                continue
            k, lab, kind, fn = items[i]
            bold = kind in ('sec', 'tot')
            fl = {'sec': FILL_SUB, 'tot': FILL_SUB}.get(kind)
            put(ws, f'{lc}{r}', lab, F_TXTB if bold else (F_NOTE if kind == 'sub' else F_TXT), fl, align=AL)
            for vc, hc in VC[side]:
                d, suf = DT[hc]
                if fn is not None:
                    f = fn(d, suf, hc)
                elif k == '其他应收':
                    f = f'SUM({vc}{R["L_个人持有"]}:{vc}{R["L_往来应收"]})'
                elif k == '其他应付':
                    f = f'SUM({vc}{R["R_个人借款"]}:{vc}{R["R_往来应付"]})'
                elif k == '负债合计':
                    f = '+'.join(f'{vc}{R["R_" + x]}' for x in ('应付账款', '预收账款', '应付工资', '其他应付', '短期借款', '应交税费'))
                elif k == '未分配':
                    f = f'{vc}{R["R_期初未分配"]}+{vc}{R["R_累计利润"]}'
                elif k == '权益合计':
                    f = f'{vc}{R["R_实收资本"]}+{vc}{R["R_未分配"]}'
                _cell(ws, f'{vc}{r}', f'={f}', F_AUTOB if bold else (F_NOTE if kind == 'sub' else F_AUTO), fl)
    put(ws, f'A{RT}', '资产合计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'D{RT}', '负债和所有者权益合计', F_TXTB, FILL_TOT, align=AL)
    for vc in 'BC':
        f = '+'.join(f'{vc}{R["L_" + x]}' for x in ('货币资金', '应收账款', '预付账款', '其他应收', '留抵', '固定资产'))
        _cell(ws, f'{vc}{RT}', f'={f}', F_AUTOB, FILL_TOT)
    for vc in 'EF':
        _cell(ws, f'{vc}{RT}', f'={vc}{R["R_负债合计"]}+{vc}{R["R_权益合计"]}', F_AUTOB, FILL_TOT)
    ws.row_dimensions[RT].height = 22

    # 核对
    rc = RT + 2
    BS_R['核对'] = rc
    put(ws, f'A{rc}', '核对：资产 − 负债 − 所有者权益（应为 0）', F_TXTB, C_LBL, align=AL)
    _cell(ws, f'B{rc}', f'=ROUND(B{RT}-E{RT},2)', F_AUTOB, FILL_AUTO)
    _cell(ws, f'C{rc}', f'=ROUND(C{RT}-F{RT},2)', F_AUTOB, FILL_AUTO)
    ws.conditional_formatting.add(f'B{rc}:C{rc}', FormulaRule(formula=[f'ROUND(B{rc},2)<>0'], font=F_RED_NOTE, fill=CF_WARN))
    ws.conditional_formatting.add(f'B{rc}:C{rc}', FormulaRule(formula=[f'ROUND(B{rc},2)=0'], font=F_OK, fill=CF_OK))
    ws.merge_cells(f'D{rc}:F{rc}')
    put(ws, f'D{rc}', (f'=IF(AND(B{rc}=0,C{rc}=0),"✓ 平了：资产＝负债＋所有者权益（年初、期末都平）",'
                       f'"✗ 不平：哪一类钱没归进来（多半是成本费用没进利润表某一行，或期初余额类型填错），去【数据校验】看")'
                       f'&IF(ROUND(K{HR["核对差额"]},2)<>0,"；成本费用有 "&TEXT(K{HR["核对差额"]},"#,##0.00")&" 元没进利润表任何一行","")'
                       f'&IF(ROUND(J{HR["期初表未分配"]},2)<>0,"；【期初余额】里填了未分配利润 "&TEXT(J{HR["期初表未分配"]},"#,##0.00")'
                       f'&"，不用它：期初未分配利润以自动倒挤为准（"&TEXT(J{HR["期初未分配"]},"#,##0.00")&"）","")'),
        F_TXTB, align=ALW, border=False)
    ws.row_dimensions[rc].height = 34
    ws.conditional_formatting.add(f'D{rc}', FormulaRule(formula=[f'OR(B{rc}<>0,C{rc}<>0)'], font=F_RED_NOTE))
    rr = rc + 1
    put(ws, f'A{rr}', '本期利润（未分配利润 期末 − 年初）', F_NOTE, align=AL)
    _cell(ws, f'C{rr}', f'=C{R["R_未分配"]}-B{R["R_未分配"]}', F_NOTE)
    ws.merge_cells(f'D{rr}:F{rr}')
    put(ws, f'D{rr}', '← 资产负债表日＝截止日时，这个数等于【利润表】的本年累计净利润', F_NOTE, align=AL, border=False)
    BS_R['本期利润'] = rr

    r = rr + 2
    _note(ws, r, 'A', last, '说明', F_TXTB)
    for text in [
        '货币资金＝银行、现金、专户（农民工专户等）的余额；老板/员工自己的微信、银行卡（个人账户）不算，按人算：结余是正的＝他手上有公司的钱（左边「个人手上的公司钱」），'
        '是负的＝他替公司垫了钱还没报销（右边「个人垫付还没报销的」）。公司还他垫付的钱记「报销还款」。',
        '应收账款＝按收入口径确认的收入 − 已收工程款（逐个项目算，收多了的项目放右边「预收账款」）；应付账款＝应付登记 − 已付（逐个供应商算，付多了的放左边「预付账款」）。',
        '应付工资＝建账前欠薪 ＋ 考勤应发（考勤起算月以后）− 已发工资（发给有考勤的人的）。过账人员（工资社保只是在公司账上过一下的）的钱在「代收代付」。',
        '应交税费＝按发票估算的税金（同利润表）− 已交的税费 ＋ 期初；是负数（交多了、进项多）放左边「留抵」。往来款按全部单位/个人的净额，正的放左边、负的放右边。',
        '未分配利润＝建账前攒下的（期初余额自动倒推：期初各账户余额＋期初余额表里的资产 − 负债 − 建账前欠薪 − 期初实收资本）＋ 建账以来每年的利润（跟利润表一个算法）。',
        '固定资产按买价记，不提折旧（管理用）。金额都含税。这是给老板看家底的管理报表，报税用的报表以代账会计的为准。',
    ]:
        r += 1
        _note(ws, r, 'A', last, text, F_NOTE, h=30)
    ws.freeze_panes = f'A{H + 1}'
    print_setup(ws, f'{H}:{H}', landscape=True)


# ═══════════════════════════ 盈亏平衡表 ═══════════════════════════
BE_R = {}
BE_PJ0 = 7                                         # 隐藏块：所选期间的项目累计收入（第 7 行起 N_PJ 行）
BE_HC = dict(名称='R', 开票0='S', 收款0='T', 产值0='U', 收入0='V', 开票1='W', 收款1='X', 产值1='Y', 收入1='Z')


def build_be(ws):
    MC = [CL(3 + i) for i in range(12)]           # C..N ＝ 1～12 月
    last = 'N'
    widths(ws, {'A': 40, 'B': 15, **{c: 12.5 for c in MC}, 'O': 9, 'P': 3})
    tip = ('💡 收入要到多少才不亏（保本点）。C3、E3 黄格填起止日期（空着＝首页报表年度的年初～截止日）；右边是报表年度每个月。'
           '变动成本＝跟着干活走的钱（材料、分包、机械、工地工人工资、其他直接费，加上按发票估算的税金）；'
           '固定费用＝不干活也要花的钱（办公室开销、管理人员工资、招待、车辆、利息手续费）。数跟【利润表】同一个算法。')
    title(ws, '盈 亏 平 衡 表（保本点）', last, C_RPT, tip)
    home_link(ws, 'O1')
    _sel_row(ws, '起')
    put(ws, 'C4', '=IF(ISNUMBER(C3),INT(C3),P_年初)', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'E4', '=IF(ISNUMBER(E3),INT(E3),P_截止)', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('F4:N4')
    put(ws, 'F4', '="收入确认口径："&P_收入口径&"；每个月＝"&P_年度&" 年各月（超过截止日的月份是 0）"', F_NOTE, align=AL, border=False)
    a0, b0 = '($C$4-1)', '$E$4'

    # ── 隐藏块：所选期间的工程收入要逐项目按口径算（跟 _汇 一样的写法）──
    C = BE_HC
    p1 = BE_PJ0 + N_PJ - 1
    rl, ra, rn, rt = p1 + 1, p1 + 2, p1 + 3, p1 + 4
    # 第 6 行：列名（R 列＝说明：0＝到「起」前一天、1＝到「止」的累计）
    for k, c in C.items():
        ws[f'{c}{BE_PJ0 - 1}'] = '（隐藏）项目' if k == '名称' else k
        ws[f'{c}{BE_PJ0 - 1}'].font = F_HELP
    for i in range(N_PJ):
        r = BE_PJ0 + i
        a = f'$R{r}'
        ws[f'R{r}'] = f'=INDEX(项目_名称,{i + 1})&""'
        for j, d in (('0', a0), ('1', b0)):
            kp, sk, cz, rv = C['开票' + j], C['收款' + j], C['产值' + j], C['收入' + j]
            ws[f'{kp}{r}'] = f'=IF({a}="",0,{cum_kp(a, d)})'
            ws[f'{sk}{r}'] = f'=IF({a}="",0,{cum_sk(a, d)})'
            ws[f'{cz}{r}'] = f'=IF({a}="",0,{cum_cz(a, d)})'
            ws[f'{rv}{r}'] = f'={rev_formula(f"{kp}{r}", f"{sk}{r}", f"{cz}{r}")}'
    ws[f'R{rl}'], ws[f'R{ra}'], ws[f'R{rn}'], ws[f'R{rt}'] = '名单合计', '全部', '未指定项目', '公司合计'
    for j, d in (('0', a0), ('1', b0)):
        kp, sk, cz, rv = C['开票' + j], C['收款' + j], C['产值' + j], C['收入' + j]
        for c in (kp, sk, cz, rv):
            ws[f'{c}{rl}'] = f'=SUM({c}{BE_PJ0}:{c}{p1})'
        ws[f'{kp}{ra}'] = f'=SUMIFS(应_开票额,应_日期,"<="&{d})'
        ws[f'{sk}{ra}'] = f'=SUMIFS(收_净额,收_归类,"工程款收款",收_日期,"<="&{d})'
        ws[f'{cz}{ra}'] = f'=SUMIFS(应_产值额,应_日期,"<="&{d})'
        for c in (kp, sk, cz):
            ws[f'{c}{rn}'] = f'=ROUND({c}{ra}-{c}{rl},2)'
        ws[f'{rv}{rn}'] = f'={rev_formula(f"{kp}{rn}", f"{sk}{rn}", f"{cz}{rn}")}'
        ws[f'{rv}{rt}'] = f'={rv}{rl}+{rv}{rn}'
    for r in range(BE_PJ0, rt + 1):
        for c in C.values():
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *C.values())
    rev_sel = f'{C["收入1"]}{rt}-{C["收入0"]}{rt}'

    H = 5
    header(ws, H, [('A', '项    目'), ('B', '所选期间')] + [(c, f'=P_年度&"年{i + 1}月"') for i, c in enumerate(MC)], C_RPT, height=30)
    cols = [('B', a0, b0, rev_sel)]
    for i, c in enumerate(MC):
        pa = '年初前' if i == 0 else f'{i}月'
        cols.append((c, f'汇_日{pa}', f'汇_日{i + 1}月', f'汇_收入{i + 1}月-汇_收入{pa}'))
    rows = [
        ('收入', '收入（工程收入，按收入确认口径）', 'sec'),
        ('变动', '变动成本（直接成本＋税金）', 'line'),
        ('直接', '　　其中：直接成本（材料、分包、机械、人工、其他直接费）', 'sub'),
        ('税金', '　　　　　税金及附加（估算）', 'sub'),
        ('边际', '边际贡献（收入 − 变动成本）', 'tot'),
        ('边际率', '边际贡献率（每收 100 元剩下多少）', 'pct'),
        ('固定', '固定费用（管理费用＋财务费用）', 'line'),
        ('经营利润', '经营利润（边际贡献 − 固定费用）', 'tot'),
        ('保本', '保本收入（固定费用 ÷ 边际贡献率）', 'grand'),
        ('安全', '安全边际（收入 − 保本收入）', 'line'),
        ('安全率', '安全边际率（安全边际 ÷ 收入）', 'pct'),
    ]
    R = BE_R
    for i, (k, _l, _t) in enumerate(rows):
        R[k] = H + 1 + i
    for k, lab, kind in rows:
        r = R[k]
        bold = kind in ('sec', 'tot', 'grand')
        fl = {'sec': FILL_SUB, 'tot': FILL_SUB, 'grand': FILL_TOT}.get(kind)
        put(ws, f'A{r}', lab, F_TXTB if bold else (F_NOTE if kind == 'sub' else F_TXT), fl, align=AL)
        for c, a, b, rev in cols:
            L = lines(a, b)
            g = lambda x: f'{c}{R[x]}'
            if k == '收入':
                f = rev
            elif k == '变动':
                f = f'{g("直接")}+{g("税金")}'
            elif k == '直接':
                f = direct_cost(L)
            elif k == '税金':
                f = L['税金']
            elif k == '边际':
                f = f'{g("收入")}-{g("变动")}'
            elif k == '边际率':
                f = f'IF(N({g("收入")})=0,"",{g("边际")}/{g("收入")})'
            elif k == '固定':
                f = f'({L["管理费用"]})+({L["财务费用"]})'
            elif k == '经营利润':
                f = f'{g("边际")}-{g("固定")}'
            elif k == '保本':
                f = (f'IF(N({g("收入")})=0,IF(N({g("固定")})=0,0,"没有收入"),IF({g("边际")}<=0,"现在的毛利盖不住费用",'
                     f'{g("固定")}/{g("边际率")}))')
            elif k == '安全':
                f = f'IF(ISNUMBER({g("保本")}),{g("收入")}-{g("保本")},"—")'
            elif k == '安全率':
                f = f'IF(AND(ISNUMBER({g("安全")}),N({g("收入")})<>0),{g("安全")}/{g("收入")},"—")'
            font = F_AUTOB if bold else (F_NOTE if kind == 'sub' else F_AUTO)
            _cell(ws, f'{c}{r}', f'={f}', font, fl, PCT if kind == 'pct' else MONEY)
    rb = R['保本']
    ws.conditional_formatting.add(f'B{rb}:{last}{rb}', FormulaRule(formula=[f'ISTEXT(B{rb})'], font=F_RED_NOTE))
    ws.conditional_formatting.add(f'B{R["安全"]}:{last}{R["安全"]}',
                                  FormulaRule(formula=[f'AND(ISNUMBER(B{R["安全"]}),B{R["安全"]}<0)'], fill=CF_WARN))

    # 大白话
    r = R['安全率'] + 2
    _note(ws, r, 'A', last, '大白话', F_TXTB)
    g = lambda x: f'B{R[x]}'
    r += 1
    BE_R['结论'] = r
    _note(ws, r, 'A', last,
          (f'=IF(N({g("收入")})=0,"所选期间没有确认收入，算不了保本点（看看【应收登记】开票/产值、【收支登记】工程款有没有记）。",'
           f'"所选期间确认工程收入 "&TEXT({g("收入")},"#,##0")&" 元，扣掉材料、分包、机械、工地工资、其他直接费和税金（变动成本 "'
           f'&TEXT({g("变动")},"#,##0")&" 元），剩下 "&TEXT({g("边际")},"#,##0")&" 元，也就是每收 100 元剩 "&TEXT(N({g("边际率")})*100,"0.0")'
           f'&" 元。这段时间公司的固定开销（管理费用＋财务费用）是 "&TEXT({g("固定")},"#,##0")&" 元，"'
           f'&IF(N({g("边际")})<=0,"现在的毛利盖不住费用：干得越多亏得越多，先看材料、分包、人工是不是超了，或者报价太低。",'
           f'"收入要做到 "&TEXT(N({g("保本")}),"#,##0")&" 元才保本；"&IF(N({g("安全")})>=0,"现在比保本线多 "&TEXT(N({g("安全")}),"#,##0")'
           f'&" 元（安全边际率 "&TEXT(N({g("安全率")}),"0.0%")&"），是赚钱的。","现在还差 "&TEXT(-N({g("安全")}),"#,##0")&" 元才到保本线，这段时间在亏。")))'),
          F_TXTB, h=48)
    for text in [
        '边际贡献率：每收 100 元工程款，扣掉跟着干活走的钱（变动成本）还剩多少元——这些钱先用来付固定费用，付完了剩下的才是利润。',
        '保本收入：这段时间至少要确认这么多收入才不亏（固定费用 ÷ 边际贡献率）。边际贡献是负的（干一单亏一单），就算不出保本点，显示「现在的毛利盖不住费用」。',
        '安全边际：实际收入比保本收入多出来的部分；安全边际率越高越稳，低于 10% 就要当心了，是负数说明这段时间在亏。',
        '一个月一个月看会跳得厉害（工程款不是每个月都收、票不是每个月都开），看全年或几个月合起来更准。'
        '利息收入、营业外支出、未分类收支不算在这里，所以「经营利润」跟【利润表】的净利润会差一点。',
    ]:
        r += 1
        _note(ws, r, 'A', last, text, F_NOTE, h=30)
    ws.freeze_panes = f'B{H + 1}'
    print_setup(ws, f'{H}:{H}', landscape=True)


def build(wb, ctx=None):
    build_pl(wb[SH_PL])
    build_bs(wb[SH_BS])
    build_be(wb[SH_BE])
