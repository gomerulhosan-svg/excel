# -*- coding: utf-8 -*-
"""查看用的表（续）：【工资表】（选月份，跟你原来的工资表一个样）【工资汇总】【用工成本表】【发票统计】"""
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *
import cats
import s_agg
from s_views import _sheet, _lbl, _val, jc, GREY

PAY_ROWS = 150          # 工资表一个月最多多少行
PAY_PJ = 8              # 工资表项目列
PAYS_ROWS = 200         # 工资汇总最多多少人
LAB_PJ = 14             # 用工成本表项目列
INV_UNITS = 120         # 发票统计按单位最多多少家


def _money_rows(ws, r0, r1, cols, money, extra=None, left=()):
    style_rows(ws, r0, r1, cols, auto=cols, fmts={**{c: MONEY for c in money}, **(extra or {})},
               aligns={**{c: AR for c in money}, **{c: AL for c in left}})
    for r in range(r0, r1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE


# ═══════════════════════════ 工资表（选一个月） ═══════════════════════════
def build_pay(wb, ctx):
    PD = [CL(7 + j) for j in range(PAY_PJ)]                  # G..N 各项目天数
    PA = [CL(17 + j) for j in range(PAY_PJ)]                 # Q..X 各项目小计
    (C_DAYS, C_RATE) = 'O', 'P'
    (C_BASE, C_ALW, C_ADJ, C_PAY, C_TAX, C_SOC, C_NET, C_PAID, C_OFF, C_REF, C_LEFT, C_OWE, C_SIGN, C_NOTE) = \
        [CL(25 + i) for i in range(14)]                       # Y..AL
    last = C_NOTE
    ws = _sheet(wb, SH_PAY, C_HOME, '工 资 表（选月份 · 每人在各工地的天数和工资 · 本月发了多少、还欠多少）',
                '💡 黄格子选月份（空着＝最近录了考勤的那个月）。全部从【考勤工资】自动来：每人一行，各工地天数、日工资、各工地小计，跟你原来的工资表一个样。'
                '「本月发放」＝【资金流水】这个月发给他的（类别「工资」），「总包代发」＝【代发抵账】里总包直接发给他的，「已退回」＝工人多拿了退回来的；'
                '「累计欠薪」＝到这个月底公司还欠他多少（负数＝多发了，他欠公司）。要打印发工资，筛掉空行就行。',
                last, {'A': 5, 'B': 9, 'C': 18, 'D': 12, 'E': 20, 'F': 12, **{c: 8 for c in PD}, C_DAYS: 7, C_RATE: 8,
                       **{c: 10 for c in PA}, C_BASE: 11, C_ALW: 9, C_ADJ: 9, C_PAY: 11, C_TAX: 8, C_SOC: 8, C_NET: 11,
                       C_PAID: 11, C_OFF: 10, C_REF: 9, C_LEFT: 11, C_OWE: 11, C_SIGN: 8, C_NOTE: 16})
    _lbl(ws, 'B3', '月份')
    last_att = next((a['mon'] for a in sorted(ctx['att_rows'], key=lambda a: a['mon'], reverse=True)), None)
    put(ws, 'C3', last_att, F_SEL, FILL_SEL, 'yyyy"年"m"月"', AC)
    dv_date(ws, 'C3')
    latest = f'MAX({atr(AT_MON)})'
    M = '$D$3'
    put(ws, 'D3', f'=IF(ISNUMBER($C$3),DATE(YEAR($C$3),MONTH($C$3),1),IF({latest}=0,DATE({AX_Y},1,1),DATE(YEAR({latest}),MONTH({latest}),1)))',
        F_KPI_L, FILL_AUTO, 'yyyy"年"m"月"', AC)
    YM = '$E$3'
    put(ws, 'E3', f'=YEAR({M})*100+MONTH({M})', F_HELP, border=False)
    HP, HA = 'BC', 'BD'
    cand = AX_ALL_N
    for i in range(cand):
        r = 1 + i
        c_ = f'{SH_AUX}!${AX_ALL}${r}'
        n = '+'.join(f'COUNTIFS({atr(p)},{c_},{atr(AT_YM)},{YM})' for p in AT_PJS)
        prev = f'{HP}{r - 1}+' if i else ''
        ws[f'{HP}{r}'] = f'={prev}AND({c_}<>"",({n})>0)'
        ws[f'{HP}{r}'].font = F_HELP
    na = AT_R1 - AT_R0 + 1
    counter(ws, HA, 1, na, lambda i: f'{SH_ATT}!${AT_YM}${AT_R0 + i}={YM}')
    npj = cnt(HP, 1, cand)
    H0 = 5
    R0 = H0 + 1
    heads = [('A', '序号'), ('B', '姓名'), ('C', '身份证号'), ('D', '手机号'), ('E', '银行账号'), ('F', '开户行'),
             (C_DAYS, '合计\n天数'), (C_RATE, '日工资\n(月工资)'), (C_BASE, '基本工资'), (C_ALW, '补贴'), (C_ADJ, '其他加减'),
             (C_PAY, '应发工资'), (C_TAX, '代扣\n个税'), (C_SOC, '代扣\n社保'), (C_NET, '实发应付'), (C_PAID, '本月发放'),
             (C_OFF, '总包代发'), (C_REF, '已退回'), (C_LEFT, '本月还欠'), (C_OWE, '累计欠薪\n(月底)'), (C_SIGN, '签字'), (C_NOTE, '备注')]
    header(ws, H0, heads, C_HOME, height=40)
    for j in range(PAY_PJ):
        nm = f'IFERROR(INDEX({AX_ALL_R},MATCH({j + 1},${HP}$1:${HP}${cand},0)),"")'
        if j == PAY_PJ - 1:
            nm = f'IF({npj}>{PAY_PJ},"其他项目",{nm})'
        put(ws, f'{PD[j]}{H0}', f'={nm}', F_HDR, fill('FF548235'), align=ACW)
        put(ws, f'{PA[j]}{H0}', f'=IF({PD[j]}${H0}="","",{PD[j]}${H0}&"小计")', F_HDR, fill('FF548235'), align=ACW)
    HID = ['AN', 'AO'] + [CL(CI('AP') + i) for i in range(12)]    # idx、第一次出现、4 对 项目/天数/金额
    PJc, DDc, AMc = HID[2:6], HID[6:10], HID[10:14]
    UN = lambda col, nm: f'IFERROR(INDEX({rng(SH_UNIT, col, UN_R0, UN_R1)},MATCH({nm},{UN_NAMES},0))&"","")'
    rngj = lambda b: f'{jr(J_YM)},">="&{AX_OYM},{jr(J_YM)},"<="&{b}'
    for k in range(PAY_ROWS):
        r = R0 + k
        ix = f'$AN{r}'
        ws[ix] = f'={kth(k + 1, HA, 1, na)}'
        a = lambda col: f'INDEX({atr(col)},{ix})'
        nm = f'$B{r}'
        ws[f'AO{r}'] = f'=IF({ix}=0,0,IF(COUNTIF($B${R0}:B{r},{nm})=1,1,0))'
        for p_, d_, m_, sp, sd, sa in zip(PJc, DDc, AMc, AT_PJS, AT_DDS, AT_AMTS):
            ws[f'{p_}{r}'] = f'=IF({ix}=0,"",{a(sp)}&"")'
            ws[f'{d_}{r}'] = f'=IF({ix}=0,0,N({a(sd)}))'
            ws[f'{m_}{r}'] = f'=IF({ix}=0,0,N({a(sa)}))'
        for c in HID:
            ws[f'{c}{r}'].font = F_HELP
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",{a(AT_NAME)})'
        for c, col in (('C', UN_ID), ('D', UN_TEL), ('E', UN_BANKNO), ('F', UN_BANK)):
            ws[f'{c}{r}'] = f'=IF({ix}=0,"",{UN(col, nm)})'
        for j in range(PAY_PJ):
            h = f'{PD[j]}${H0}'
            dsum = '+'.join(f'({p_}{r}={h})*{d_}{r}' for p_, d_ in zip(PJc, DDc))
            asum = '+'.join(f'({p_}{r}={h})*{m_}{r}' for p_, m_ in zip(PJc, AMc))
            if j == PAY_PJ - 1:
                prevd = '+'.join(f'N({PD[t]}{r})' for t in range(PAY_PJ - 1))
                preva = '+'.join(f'N({PA[t]}{r})' for t in range(PAY_PJ - 1))
                ws[f'{PD[j]}{r}'] = (f'=IF(OR({ix}=0,{h}=""),"",IF({h}="其他项目",'
                                     f'{"+".join(d_ + str(r) for d_ in DDc)}-({prevd}),{dsum}))')
                ws[f'{PA[j]}{r}'] = (f'=IF(OR({ix}=0,{h}=""),"",IF({h}="其他项目",'
                                     f'{"+".join(m_ + str(r) for m_ in AMc)}-({preva}),{asum}))')
            else:
                ws[f'{PD[j]}{r}'] = f'=IF(OR({ix}=0,{h}=""),"",{dsum})'
                ws[f'{PA[j]}{r}'] = f'=IF(OR({ix}=0,{h}=""),"",{asum})'
        for c, col in ((C_DAYS, AT_DAYS), (C_RATE, AT_RATE), (C_BASE, AT_BASE), (C_ALW, AT_ALLOW), (C_ADJ, AT_ADJ), (C_PAY, AT_PAY),
                       (C_TAX, AT_TAX), (C_SOC, AT_SOC), (C_NET, AT_NETPAY)):
            ws[f'{c}{r}'] = f'=IF({ix}=0,"",N({a(col)}))'
        ws[f'{C_NOTE}{r}'] = f'=IF({ix}=0,"",{a(AT_NOTE)}&"")'
        first = f'$AO{r}=1'
        mrng = f'{jr(J_YM)},{YM}'
        ws[f'{C_PAID}{r}'] = (f'=IF({ix}=0,"",IF({first},-SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"应付职工薪酬",'
                              f'{jr(J_NET)},"<0",{mrng}),0))')
        ws[f'{C_OFF}{r}'] = (f'=IF({ix}=0,"",IF({first},SUMIFS({ofr(OF_AMT)},{ofr(OF_WHO)},{nm},{ofr(OF_TYPE)},"总包代发工资",'
                             f'{ofr(OF_YM)},{YM}),0))')
        ws[f'{C_REF}{r}'] = (f'=IF({ix}=0,"",IF({first},SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"应付职工薪酬",'
                             f'{jr(J_NET)},">0",{mrng}),0))')
        ws[f'{C_LEFT}{r}'] = f'=IF({ix}=0,"",{C_NET}{r}-{C_PAID}{r}-{C_OFF}{r}+{C_REF}{r})'
        owe = (f'N(IFERROR(INDEX({rng(SH_UNIT, UN_OWE0, UN_R0, UN_R1)},MATCH({nm},{UN_NAMES},0)),0))+SUMIFS({atr(AT_NETPAY)},{atr(AT_NAME)},{nm},{atr(AT_YM)},">="&{AX_OYM},{atr(AT_YM)},"<="&{YM})'
               f'+SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"应付职工薪酬",{rngj(YM)})'
               f'-SUMIFS({ofr(OF_AMT)},{ofr(OF_WHO)},{nm},{ofr(OF_TYPE)},"总包代发工资",{ofr(OF_YM)},">="&{AX_OYM},{ofr(OF_YM)},"<="&{YM})')
        ws[f'{C_OWE}{r}'] = f'=IF({ix}=0,"",IF({first},ROUND({owe},2),""))'
    R1 = R0 + PAY_ROWS - 1
    cols = [CL(i) for i in range(1, CI(last) + 1)]
    money = PA + [C_BASE, C_ALW, C_ADJ, C_PAY, C_TAX, C_SOC, C_NET, C_PAID, C_OFF, C_REF, C_LEFT, C_OWE]
    _money_rows(ws, R0, R1, cols, money, {**{c: '0.##;-0.##;""' for c in PD + [C_DAYS]}, C_RATE: '#,##0.##', 'E': '@'},
                left=('C', 'E', 'F', C_NOTE))
    ws.conditional_formatting.add(f'{C_OWE}{R0}:{C_OWE}{R1}', FormulaRule(formula=[f'N(${C_OWE}{R0})<0'], fill=fill('FFFFEB9C')))
    put(ws, 'B4', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in cols:
        if c in ('A', 'B', 'C', 'D', 'E', 'F', C_RATE, C_SIGN, C_NOTE):
            put(ws, f'{c}4', None, fill_=FILL_TOT)
            continue
        fmt = '0.##;-0.##;""' if c in PD + [C_DAYS] else MONEY
        put(ws, f'{c}4', f'=SUM({c}{R0}:{c}{R1})', F_AUTOB, FILL_TOT, fmt, AR)
    put(ws, 'F3', f'=IF({cnt(HA, 1, na)}>{PAY_ROWS},"⚠ 这个月有"&{cnt(HA, 1, na)}&"行考勤，只列前{PAY_ROWS}行","")'
                  f'&IF({npj}>{PAY_PJ},"  这个月跑了"&{npj}&"个项目，第{PAY_PJ}列起合在「其他项目」","")', F_RED, border=False)
    ws.merge_cells(f'F3:{PA[3]}3')
    hide(ws, *HID, HP, HA)
    ws.freeze_panes = f'C{R0}'
    ws.auto_filter.ref = f'A{H0}:{last}{R1}'
    return ws


# ═══════════════════════════ 工资汇总（每人一行：本年应发、已发、欠薪） ═══════════════════════════
def build_pays(wb, ctx):
    MC = [CL(18 + i) for i in range(12)]          # R..AC 各月应发
    last = 'AD'
    ws = _sheet(wb, SH_PAYS, C_HOME, '工 资 汇 总（每人一行：本年应发、发了多少、还欠多少）',
                '💡 全自动，年度跟首页走。年初欠薪＋本年实发应付−本年发放−总包代发＋已退回＝剩余工资（负数＝多发了，他欠公司，下个月从工资里扣或让他退）。'
                '右边是各月应发。月薪的人某个月没录考勤，最后一列会提醒（漏录了工资就进不了成本）。',
                last, {'A': 5, 'B': 10, 'C': 9, 'D': 10, 'E': 11, 'F': 8, 'G': 12, 'H': 9, 'I': 9, 'J': 9, 'K': 12, 'L': 12,
                       'M': 11, 'N': 10, 'O': 12, 'P': 14, 'Q': 2, **{c: 9 for c in MC}, 'AD': 20})
    put(ws, 'B3', f'={AX_Y}&"年  截止 "&TEXT({AX_E},"m月d日")', F_KPI_L, align=AL, border=False)
    ws.merge_cells('B3:D3')
    nu = UN_R1 - UN_R0 + 1
    HC = 'BA'

    def cond(i):
        b = BX_R0 + i
        t = f'{SH_BALX}!${BX_TYPE}${b}'
        return (f'AND(OR({t}="管理人员",{t}="工人"),OR({SH_BALX}!${BX_NETPAY}${b}<>0,{SH_BALX}!${BX_WPAID}${b}<>0,'
                f'{SH_BALX}!${BX_WG_E}${b}<>0,{SH_BALX}!${BX_OFFW}${b}<>0))')
    counter(ws, HC, 1, nu, cond)
    H0 = 5
    R0 = H0 + 1
    header(ws, H0, [('A', '序号'), ('B', '姓名'), ('C', '类型'), ('D', '现在的\n日薪/月薪'), ('E', '年初欠薪'), ('F', '本年\n天数'),
                    ('G', '本年应发'), ('H', '其中补贴'), ('I', '代扣个税'), ('J', '代扣社保'), ('K', '本年\n实发应付'), ('L', '本年发放'),
                    ('M', '总包代发'), ('N', '已退回'), ('O', '剩余工资\n（欠薪）'), ('P', '说明'), ('AD', '提醒')], C_HOME, height=40)
    for i, c in enumerate(MC):
        put(ws, f'{c}{H0}', f'{i + 1}月', F_HDR, fill('FF548235'), align=AC)
    y0, y1 = f'MAX({AX_YM0},{AX_OYM})', AX_YM1
    ya = f'{atr(AT_YM)},">="&{y0},{atr(AT_YM)},"<="&{y1}'
    yj = f'{jr(J_YM)},">="&{y0},{jr(J_YM)},"<="&{y1}'
    yo = f'{ofr(OF_YM)},">="&{y0},{ofr(OF_YM)},"<="&{y1}'
    for k in range(PAYS_ROWS):
        r = R0 + k
        ix = f'$AZ{r}'
        ws[ix] = f'={kth(k + 1, HC, 1, nu)}'
        ws[ix].font = F_HELP
        nm = f'$B{r}'
        b = lambda col: f'INDEX({s_agg.bx(col)},{ix})'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX({UN_NAMES},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX({UN_TYPES_R},{ix}))'
        eff = f'_xlfn.MAXIFS({RT_DATES},{RT_NAMES},{nm},{RT_DATES},"<="&{AX_E})'
        ws[f'D{r}'] = (f'=IF({ix}=0,"",IF({eff}=0,"",SUMIFS({RT_DAYS},{RT_NAMES},{nm},{RT_DATES},{eff})'
                       f'+SUMIFS({RT_MONS},{RT_NAMES},{nm},{RT_DATES},{eff})))')
        ws[f'E{r}'] = f'=IF({ix}=0,"",{b(BX_WG_B)})'
        sa = lambda col: f'SUMIFS({atr(col)},{atr(AT_NAME)},{nm},{ya})'
        ws[f'F{r}'] = f'=IF({ix}=0,"",{sa(AT_DAYS)})'
        ws[f'G{r}'] = f'=IF({ix}=0,"",{sa(AT_PAY)})'
        ws[f'H{r}'] = f'=IF({ix}=0,"",{sa(AT_ALLOW)})'
        ws[f'I{r}'] = f'=IF({ix}=0,"",{sa(AT_TAX)})'
        ws[f'J{r}'] = f'=IF({ix}=0,"",{sa(AT_SOC)})'
        ws[f'K{r}'] = f'=IF({ix}=0,"",{sa(AT_NETPAY)})'
        ws[f'L{r}'] = f'=IF({ix}=0,"",-SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"应付职工薪酬",{jr(J_NET)},"<0",{yj}))'
        ws[f'M{r}'] = f'=IF({ix}=0,"",SUMIFS({ofr(OF_AMT)},{ofr(OF_WHO)},{nm},{ofr(OF_TYPE)},"总包代发工资",{yo}))'
        ws[f'N{r}'] = f'=IF({ix}=0,"",SUMIFS({jr(J_NET)},{jr(J_UN)},{nm},{jr(J_LINE)},"应付职工薪酬",{jr(J_NET)},">0",{yj}))'
        ws[f'O{r}'] = f'=IF({ix}=0,"",{b(BX_WG_E)})'
        ws[f'P{r}'] = (f'=IF({ix}=0,"",IF(ROUND(E{r}+K{r}-L{r}-M{r}+N{r}-O{r},2)<>0,"✗ 对不上，查一下",'
                       f'IF(O{r}<0,"多发了，他欠公司",IF(O{r}>0,"公司还欠他",""))))')
        for i, c in enumerate(MC):
            ym = f'{AX_Y}*100+{i + 1}'
            ws[f'{c}{r}'] = f'=IF({ix}=0,"",SUMIFS({atr(AT_PAY)},{atr(AT_NAME)},{nm},{atr(AT_YM)},{ym}))'
        first = f'MINIFS({atr(AT_YM)},{atr(AT_NAME)},{nm},{atr(AT_YM)},">="&{y0})'
        out_ = f'{SH_UNIT}!${UN_OUT}$4:${UN_OUT}${UN_R1}'
        outd = f'INDEX({out_},{ix})'
        miss = '+'.join(f'AND({AX_Y}*100+{i + 1}>=_xlfn.{first},{AX_Y}*100+{i + 1}<={y1},N({MC[i]}{r})=0,'
                        f'OR(NOT(ISNUMBER({outd})),DATE({AX_Y},{i + 1},1)<={outd}))' for i in range(12))
        ws[f'AD{r}'] = (f'=IF({ix}=0,"",IF(AND(C{r}="管理人员",D{r}<>"",COUNTIFS({RT_NAMES},{nm},{RT_MONS},">0")>0),'
                        f'IF(({miss})>0,"⚠ 有"&({miss})&"个月没录考勤",""),""))')
    R1 = R0 + PAYS_ROWS - 1
    cols = [CL(i) for i in range(1, 17)] + MC + ['AD']
    money = list('EGHIJKLMNO') + MC
    _money_rows(ws, R0, R1, cols, money, {'D': '#,##0.##', 'F': '0.##;-0.##;""'}, left=('P', 'AD'))
    ws.conditional_formatting.add(f'O{R0}:O{R1}', FormulaRule(formula=[f'N($O{R0})<0'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'P{R0}:P{R1}', FormulaRule(formula=[f'LEFT($P{R0},1)="✗"'], fill=FILL_WARN))
    ws.conditional_formatting.add(f'AD{R0}:AD{R1}', FormulaRule(formula=[f'$AD{R0}<>""'], font=F_RED))
    put(ws, 'B4', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in money + ['F']:
        put(ws, f'{c}4', f'=SUM({c}{R0}:{c}{R1})', F_AUTOB, FILL_TOT, '0.##' if c == 'F' else MONEY, AR)
    for c in 'ACDP':
        put(ws, f'{c}4', None, fill_=FILL_TOT)
    put(ws, 'AD4', f'=IF(COUNTIF(AD{R0}:AD{R1},"⚠*")>0,COUNTIF(AD{R0}:AD{R1},"⚠*")&"个管理人员有月份没录考勤","")', F_RED, FILL_TOT)
    put(ws, 'E3', f'=IF({cnt(HC, 1, nu)}>{PAYS_ROWS},"⚠ 有"&{cnt(HC, 1, nu)}&"人，只列前{PAYS_ROWS}人","")', F_RED, border=False)
    hide(ws, 'AZ', HC)
    ws.freeze_panes = f'C{R0}'
    ws.auto_filter.ref = f'A{H0}:{last}{R1}'
    return ws


# ═══════════════════════════ 用工成本表（项目 × 月，跟你原来的一个样） ═══════════════════════════
def build_lab(wb, ctx):
    PC = [CL(3 + j) for j in range(LAB_PJ)]          # C.. 各项目
    c_oth, c_sum, c_co, c_all = [CL(3 + LAB_PJ + i) for i in range(4)]
    ws = _sheet(wb, SH_LAB, C_HOME, '用 工 成 本 表（各项目每月人工 · 不含分包）',
                '💡 全自动：每个项目每个月的人工＝【考勤工资】按天数拆到这个项目的应发工资＋现结的临时工（流水类别「临时工工资」）。'
                '「以前累计」＝建账前（【期初余额】①人工）＋以前年度。「公司管理」＝管理人员在公司那部分工资（算管理费，年底按比例摊到项目）。'
                '跟你原来的用工成本表一样看：一行一个月，一列一个项目。',
                c_all, {'A': 5, 'B': 16, **{c: 12 for c in PC}, c_oth: 12, c_sum: 13, c_co: 12, c_all: 13})
    put(ws, 'B3', f'={AX_Y}&"年  截止 "&TEXT({AX_E},"m月d日")', F_KPI_L, align=AL, border=False)
    npj = PJ_R1 - PJ_R0 + 1
    HC = 'BA'

    def cum(i):
        k = PS_R0 + i
        return (f'SUMIFS({opr(OPJ_LAB)},{opr(OPJ_PJ)},{SH_PS}!${PS_NAME}${k})+{SH_PS}!${PS_COL[("人工", "前")]}${k}'
                f'+{SH_PS}!${PS_COL[("人工", "本")]}${k}')
    counter(ws, HC, 1, npj, lambda i: f'AND({SH_PS}!${PS_NAME}${PS_R0 + i}<>"",{cum(i)}<>0)')
    H0 = 5
    header(ws, H0, [('A', '序号'), ('B', '项目名称\n月份'), (c_oth, '其他项目\n/没选项目'), (c_sum, '项目合计'), (c_co, '公司管理\n(管理人员)'),
                    (c_all, '总计')], C_HOME, height=40)
    for j, c in enumerate(PC):
        put(ws, f'{c}{H0}', f'=IFERROR(INDEX({SH_PS}!${PS_NAME}${PS_R0}:${PS_NAME}${PS_R0 + npj - 1},'
                            f'MATCH({j + 1},${HC}$1:${HC}${npj},0)),"")', F_HDR, fill(C_HOME), align=ACW)
    ix = lambda c: f'IFERROR(MATCH({c}${H0},{SH_PS}!${PS_NAME}${PS_R0}:${PS_NAME}${PS_R0 + npj - 1},0),0)'
    rows = [('以前累计', None)] + [(f'{m}月', m) for m in range(1, 13)]
    R0 = H0 + 1
    for t, (lab, m) in enumerate(rows):
        r = R0 + t
        put(ws, f'A{r}', t + 1, F_AUTO, align=AC)
        put(ws, f'B{r}', lab if m is None else f'={AX_Y}&"年{m}月"', F_TXTB if m is None else F_TXT, align=AC)
        for c in PC:
            if m is None:
                v = (f'SUMIFS({opr(OPJ_LAB)},{opr(OPJ_PJ)},{c}${H0})+INDEX({s_agg.ps("人工", "前")},{ix(c)})')
            else:
                v = f'INDEX({SH_PS}!${PS_LABM[m - 1]}${PS_R0}:${PS_LABM[m - 1]}${PS_R0 + npj - 1},{ix(c)})'
            put(ws, f'{c}{r}', f'=IF({c}${H0}="","",{v})', F_AUTO, fmt=MONEY, align=AR)
        shown = '+'.join(f'N({c}{r})' for c in PC)
        if m is None:
            allp = f'SUM({opr(OPJ_LAB)})+{s_agg.ms("人工费", MS_PREV)}'
            co = s_agg.ms('管理人员工资', MS_PREV)
        else:
            allp = s_agg.ms('人工费', MS_MCOLS[m - 1])
            co = s_agg.ms('管理人员工资', MS_MCOLS[m - 1])
            allp = f'IF({AX_Y}*100+{m}>{AX_YM1},0,{allp})'
            co = f'IF({AX_Y}*100+{m}>{AX_YM1},0,{co})'
        put(ws, f'{c_oth}{r}', f'=ROUND({allp}-({shown}),2)', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'{c_sum}{r}', f'={shown}+N({c_oth}{r})', F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'{c_co}{r}', f'={co}', F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'{c_all}{r}', f'={c_sum}{r}+{c_co}{r}', F_AUTOB, fmt=MONEY, align=AR)
    r1 = R0 + 12
    for lab, f0 in (('本年小计', lambda c: f'SUM({c}{R0 + 1}:{c}{R0 + 12})'), ('建账以来合计', lambda c: f'SUM({c}{R0}:{c}{R0 + 12})')):
        r1 += 1
        put(ws, f'B{r1}', lab, F_TXTB, FILL_TOT, align=AC)
        put(ws, f'A{r1}', None, fill_=FILL_TOT)
        for c in PC + [c_oth, c_sum, c_co, c_all]:
            put(ws, f'{c}{r1}', f'={f0(c)}', F_AUTOB, FILL_TOT, MONEY, AR)
    ws.conditional_formatting.add(f'A{R0 + 1}:{c_all}{R0 + 12}', FormulaRule(formula=[f'{AX_Y}*100+ROW()-{R0}>{AX_YM1}'], font=GREY))
    put(ws, f'B{r1 + 2}', '「其他项目/没选项目」：超过 14 个项目的放这里，还有流水里临时工工资没选项目的（去【资金流水】看 ✗）。', F_NOTE, border=False)
    put(ws, 'D3', f'=IF({cnt(HC, 1, npj)}>{LAB_PJ},"⚠ 有"&{cnt(HC, 1, npj)}&"个项目有人工，前{LAB_PJ}个单列","")', F_RED, border=False)
    hide(ws, HC)
    ws.freeze_panes = f'C{R0}'
    return ws


# ═══════════════════════════ 发票统计 ═══════════════════════════
def build_invs(wb, ctx):
    ws = _sheet(wb, SH_INVS, C_INV, '发 票 统 计（按项目 · 按月 · 按单位：开了多少、收了多少、还差多少票）',
                '💡 全自动，从【发票登记】来（电子税务局导出粘进去就行），年度跟首页走。① 按项目：确认的产值开了多少票、还没开多少；'
                '材料分包机械的成本收了多少票、还差多少票（差票多了，增值税和所得税都要多交）。② 按月：销项、进项、增值税估算。'
                '③ 按单位：每家应付多少、收了多少票、欠多少票，找他们要票用。',
                'P', {'A': 5, 'B': 16, 'C': 22, 'D': 14, 'E': 14, 'F': 13, 'G': 13, 'H': 12, 'I': 13, 'J': 12, 'K': 12, 'L': 14,
                      'M': 14, 'N': 13, 'O': 9, 'P': 20})
    put(ws, 'B3', f'={AX_Y}&"年  截止 "&TEXT({AX_E},"m月d日")', F_KPI_L, align=AL, border=False)
    npj = PJ_R1 - PJ_R0 + 1
    y0, y1 = f'MAX({AX_YM0},{AX_OYM})', AX_YM1
    iy = f'{ivr(IV_YM)},">="&{y0},{ivr(IV_YM)},"<="&{y1}'
    ia = f'{ivr(IV_YM)},">="&{AX_OYM},{ivr(IV_YM)},"<="&{y1}'
    section(ws, 4, 'A', 'P', '① 按项目', C_INV)
    H0 = 5
    header(ws, H0, [('A', '序号'), ('B', '项目'), ('C', '甲方 / 总包'), ('D', '累计确认产值'), ('E', '累计开票'), ('F', '还没开票\n（确认−开票）'),
                    ('G', '本年开票\n（价税合计）'), ('H', '本年销项税'), ('I', '本年收票\n（进项）'), ('J', '本年进项\n专票税额'),
                    ('K', '本年增值税\n估算'), ('L', '累计材料分包\n机械其他成本'), ('M', '累计收票'), ('N', '还差多少票'), ('O', '收票\n比例'),
                    ('P', '提示')], C_INV, height=44)
    R0 = H0 + 2
    for i in range(npj):
        r, p, k = R0 + i, PJ_R0 + i, PS_R0 + i
        nm = f'$B{r}'
        pc = lambda m, per: f'{SH_PS}!${PS_COL[(m, per)]}${k}'
        o = lambda col: f'SUMIFS({opr(col)},{opr(OPJ_PJ)},{nm})'
        ws[f'A{r}'] = f'=IF({nm}="","",{i + 1})'
        ws[f'B{r}'] = f'={SH_PROJ}!${PJ_NAME}${p}&""'
        ws[f'C{r}'] = f'={SH_PROJ}!${PJ_CUS}${p}&""'
        ws[f'D{r}'] = f'=IF({nm}="","",{SH_AR}!$F${6 + i})'
        ws[f'E{r}'] = f'=IF({nm}="","",{SH_AR}!$J${6 + i})'
        ws[f'F{r}'] = f'=IF({nm}="","",D{r}-E{r})'
        ws[f'G{r}'] = f'=IF({nm}="","",{pc("开票", "本")})'
        ws[f'H{r}'] = f'=IF({nm}="","",SUMIFS({ivr(IV_TAXU)},{ivr(IV_DIR)},"销项",{ivr(IV_PJ)},{nm},{iy}))'
        ws[f'I{r}'] = f'=IF({nm}="","",SUMIFS({ivr(IV_TOTU)},{ivr(IV_DIR)},"进项",{ivr(IV_PJ)},{nm},{iy}))'
        ws[f'J{r}'] = f'=IF({nm}="","",{pc("进项税", "本")})'
        ws[f'K{r}'] = f'=IF({nm}="","",H{r}-J{r})'
        cost = '+'.join(f'{pc(m, per)}' for m in ('材料', '分包', '机械', '其他直接') for per in ('前', '本'))
        ws[f'L{r}'] = f'=IF({nm}="","",{o(OPJ_MAT)}+{o(OPJ_SUB)}+{o(OPJ_MACH)}+{o(OPJ_OTH)}+{cost})'
        ws[f'M{r}'] = (f'=IF({nm}="","",SUMIFS({oapr(OAP_INV)},{oapr(OAP_PJ)},{nm})'
                       f'+SUMIFS({ivr(IV_TOTU)},{ivr(IV_DIR)},"进项",{ivr(IV_PJ)},{nm},{ia}))')
        ws[f'N{r}'] = f'=IF({nm}="","",L{r}-M{r})'
        ws[f'O{r}'] = f'=IF(OR({nm}="",N(L{r})=0),"",M{r}/L{r})'
        ws[f'P{r}'] = f'=IF({nm}="","",IF(N(F{r})<0,"开票比确认的多（补确认产值？）",IF(N(O{r})<0.5,IF(N(L{r})>0,"收票不到一半",""),"")))'
    R1 = R0 + npj - 1
    _money_rows(ws, R0, R1, list('ABCDEFGHIJKLMNOP'), list('DEFGHIJKLMN'), {'O': PCT}, left=('B', 'C', 'P'))
    r = R0 - 1
    put(ws, f'B{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'DEFGHIJKLMN':
        put(ws, f'{c}{r}', f'=SUM({c}{R0}:{c}{R1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'O{r}', f'=IF(N(L{r})=0,"",M{r}/L{r})', F_AUTOB, FILL_TOT, PCT, AC)
    for c in 'ACP':
        put(ws, f'{c}{r}', None, fill_=FILL_TOT)
    # ② 按月
    M0 = R1 + 3
    section(ws, M0, 'A', 'P', '② 按月（本年）', C_INV)
    header(ws, M0 + 1, [('A', '月'), ('B', '开票张数'), ('C', '销项金额\n（不含税）'), ('D', '销项税额'), ('E', '销项价税合计'), ('F', '收票张数'),
                        ('G', '进项价税合计'), ('H', '进项专票\n税额'), ('I', '增值税估算\n（销项−进项）'), ('J', '附加税估算'),
                        ('K', '本年累计\n增值税'), ('L', '确认产值\n（不含税）'), ('M', '税负率\n（增值税÷产值）')], C_INV, height=44)
    for m in range(1, 13):
        r = M0 + 1 + m
        ym = f'{AX_Y}*100+{m}'
        crit = lambda d: f'{ivr(IV_DIR)},"{d}",{ivr(IV_YM)},{ym}'
        put(ws, f'A{r}', f'{m}月', F_TXT, align=AC)
        ws[f'B{r}'] = f'=COUNTIFS({crit("销项")})'
        ws[f'D{r}'] = f'=SUMIFS({ivr(IV_TAXU)},{crit("销项")})'
        ws[f'E{r}'] = f'=SUMIFS({ivr(IV_TOTU)},{crit("销项")})'
        ws[f'C{r}'] = f'=E{r}-D{r}'
        ws[f'F{r}'] = f'=COUNTIFS({crit("进项")})'
        ws[f'G{r}'] = f'=SUMIFS({ivr(IV_TOTU)},{crit("进项")})'
        ws[f'H{r}'] = f'=SUMIFS({ivr(IV_TAXU)},{crit("进项")},{ivr(IV_SPEC)},1)'
        ws[f'I{r}'] = f'=D{r}-H{r}'
        ws[f'J{r}'] = f'=ROUND(MAX(I{r},0)*{PA_SURR},2)'
        ws[f'K{r}'] = f'=SUM(I${M0 + 2}:I{r})'
        ws[f'L{r}'] = f'=ROUND({s_agg.ms("营业收入", MS_MCOLS[m - 1])}-{calc_vat_month(m)},2)'
        ws[f'M{r}'] = f'=IF(N(L{r})=0,"",I{r}/L{r})'
    MR1 = M0 + 13
    _money_rows(ws, M0 + 2, MR1, list('ABCDEFGHIJKLM'), list('CDEGHIJKL'), {'B': INT, 'F': INT, 'M': PCT})
    r = MR1 + 1
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'BCDEFGHIJL':
        put(ws, f'{c}{r}', f'=SUM({c}{M0 + 2}:{c}{MR1})', F_AUTOB, FILL_TOT, INT if c in 'BF' else MONEY, AR)
    put(ws, f'M{r}', f'=IF(N(L{r})=0,"",I{r}/L{r})', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, f'K{r}', None, fill_=FILL_TOT)
    ws.conditional_formatting.add(f'A{M0 + 2}:M{MR1}', FormulaRule(formula=[f'{AX_Y}*100+ROW()-{M0 + 1}>{AX_YM1}'], font=GREY))
    # ③ 按单位
    U0 = r + 3
    section(ws, U0, 'A', 'P', '③ 按单位（材料商、分包、机械：应付、已付、收票、欠票）', C_INV)
    header(ws, U0 + 1, [('A', '序号'), ('B', '单位'), ('C', '类型'), ('D', '累计应付'), ('E', '累计已付'), ('F', '累计收票'),
                        ('G', '欠票\n（应付−收票）'), ('H', '已付没给票\n（已付−收票）'), ('I', '收票比例'), ('J', '本年收票')], C_INV, height=44)
    nu = UN_R1 - UN_R0 + 1
    HC = 'BA'

    def cond(i):
        b = BX_R0 + i
        return f'OR({SH_BALX}!${BX_APAMT}${b}<>0,{SH_BALX}!${BX_INVIN}${b}<>0)'
    counter(ws, HC, 1, nu, cond)
    for k in range(INV_UNITS):
        r = U0 + 2 + k
        ix = f'$AZ{r}'
        ws[ix] = f'={kth(k + 1, HC, 1, nu)}'
        ws[ix].font = F_HELP
        b = lambda col: f'INDEX({s_agg.bx(col)},{ix})'
        nm = f'$B{r}'
        ws[f'A{r}'] = f'=IF({ix}=0,"",{k + 1})'
        ws[f'B{r}'] = f'=IF({ix}=0,"",INDEX({UN_NAMES},{ix}))'
        ws[f'C{r}'] = f'=IF({ix}=0,"",INDEX({UN_TYPES_R},{ix}))'
        ws[f'D{r}'] = f'=IF({ix}=0,"",{b(BX_APAMT)})'
        ws[f'E{r}'] = f'=IF({ix}=0,"",{b(BX_PAID)})'
        ws[f'F{r}'] = f'=IF({ix}=0,"",{b(BX_INVIN)})'
        ws[f'G{r}'] = f'=IF({ix}=0,"",D{r}-F{r})'
        ws[f'H{r}'] = f'=IF({ix}=0,"",MAX(E{r}-F{r},0))'
        ws[f'I{r}'] = f'=IF(OR({ix}=0,N(D{r})=0),"",F{r}/D{r})'
        ws[f'J{r}'] = f'=IF({ix}=0,"",SUMIFS({ivr(IV_TOTU)},{ivr(IV_DIR)},"进项",{ivr(IV_UNIT)},{nm},{iy}))'
    UR1 = U0 + 1 + INV_UNITS
    _money_rows(ws, U0 + 2, UR1, list('ABCDEFGHIJ'), list('DEFGHJ'), {'I': PCT}, left=('B',))
    ws.conditional_formatting.add(f'H{U0 + 2}:H{UR1}', FormulaRule(formula=[f'N($H{U0 + 2})>0'], font=F_RED))
    hide(ws, 'AZ', HC)
    ws.freeze_panes = f'C{R0}'
    return ws


def calc_vat_month(m):
    ym = f'{AX_Y}*100+{m}'
    return f'SUMIFS({rvr(RV_VAT)},{rvr(RV_YM)},{ym})'


def build_all(wb, ctx):
    build_pay(wb, ctx)
    build_pays(wb, ctx)
    build_lab(wb, ctx)
    build_invs(wb, ctx)
