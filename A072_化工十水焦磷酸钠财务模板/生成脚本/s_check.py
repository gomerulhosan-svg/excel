# -*- coding: utf-8 -*-
"""A072 化工 · 【数据校验】（查看表）：照 A069 s_check 的版面（分组、每项 结果 √/✗/⚠/ℹ、数量、说明「去哪改」、点名字跳过去）。
   所有判断只读隐藏层（_资 _进 _销 _产 _工 _票 _盘 _手 _位 _固 _卡 _表 _科 _物）、引擎（_料 _品 _本 _息）、_分/_凭/_矩 和报表的核对格；
   录入表只数 4 样：备注是「示例」的行、右边「这一行的问题」有字的行（跟隐藏层有内容的行数比＝右边没公式）、容量下一行起有没有东西、
   表头区 COLUMNS（插了列/删了列）。

① 各录入表（8 张）：✗ 行、⚠ 行（都不含示例行）、示例行（另数，ℹ）、右边没公式、容量外有内容、插了列；第一处 ✗ 的行号（点了跳过去）。
② 基础资料：建账日、②垫付/备用金期初、⑦ 不算的行、⑧ 不算的金额、⑧ 预付 1123 没按月摊的行（⚠）、期初试算、倒挤的期初未分配利润（ℹ）、物料期初有量没价、
   名称重复（往来单位、物料产品、人员、账户/类别/报销项目/票种）、会计科目少固定科目/编码有问题、往来单位/固定资产/固定支出的 ✗ ⚠、
   基础表超容量、插了列、示例（ℹ）。
③ 账（截止日那个月月末，P_截止月）：矩阵每月借＝贷、每张凭证借＝贷（月末凭证逐张、资金/进货/销售每 50 张一段、手工整块）、
   科目余额表 F3、资产负债表_平衡/_往来核对、利润表_核对、现金流量表_核对、Σ逐家＝1122/2202、1403＝_料、1405＝_品、4001＝在产、
   4101＝池留存、1601/1602＝卡片、122104/122105＝0、出借人利息付多了、上级科目直接记账、每月凭证张数、凭证汇总 K3、
   4101 明细/费用明细超槽位、年末结转各块核对。
④ 业务提醒：负库存、工资表没登、待定价进货/发货、待补票、垫付编号空/重复、固定支出到期没付、补票不抵税（ℹ）、
   现金账户余额负、首页截止日期比最后一笔早、日期打错的嫌疑。
顶上：C4 ✗ 合计（名称 校验_错数）、F4 ⚠ 合计（校验_警数）；ℹ 不算。另给首页用的名称：校验_负库存、校验_工资没登、校验_待定价进货笔数、
   校验_待定价进货金额、校验_待补票笔数、校验_待补票金额、校验_固定支出没付条数、校验_固定支出没付金额。

✗ 合计＝① 各表 ✗ 行数之和 ＋ 容量外有内容、插了列的表数 ＋ ②③④ 结果是 ✗ 的项数；⚠ 合计＝① ⚠ 行数之和 ＋ 右边没公式的表数 ＋ ②③④ ⚠ 项数。
隐藏辅助列从 M 列起（第 1 行是标签，第 i 个在第 i+1 行），位置由 HELPERS 顺序决定（model_check 用）。"""
import math
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font
from common import *
from layout import *                      # layout 的 norm/esc 盖过 common 的

CHK_X, CHK_W = 'C4', 'F4'
HOME_C = f'{SH_HOME}!${HOME_END[0]}${HOME_END[1:]}'
D = 'P_截止'
ME = 'DATE(P_年度,P_截止月+1,0)'          # 截止日那个月的月末
CHUNK = 50                                # 资金/进货/销售凭证每 50 张一段查借贷
F_LINK = Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single')
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC65911')
F_INFO = Font(name=YH, sz=10, bold=True, color='FF2F75B5')
F_OK = Font(name=YH, sz=10, bold=True, color='FF548235')
LAST = 'K'
SEC_C1, SEC_C2 = 'A', LAST

# ① 录入表：(表, 隐藏前缀, 备注列, 「这一行的问题」列, 最后一列（表头区）, 容量, 录入列最后一列, 常见原因)
INPUTS = [
    (SH_CASH, '资', CASH_COLS['备注'], CASH_SHOW['校验'], CASH_LAST, N_CASH,
     '没填日期或金额、日期写成 8.15、账户或类别不在【基础资料】里、收付货款/借款没选往来单位、员工报销没填报销人和报销项目、'
     '缴税没选是哪种税、垫付账户记了内部转账；⚠ 多半是垫付没填编号、付货款选错类别、日期没排序'),
    (SH_PUR, '进', PUR_COLS['备注'], PUR_SHOW['校验'], PUR_LAST, N_PUR,
     '物料、供应商不在档案里，日期早于建账日或不在本年，数量单价金额不是数字，有运费没物流单位；⚠ 还没定价（按参考单价暂估）'),
    (SH_SALE, '销', SALE_COLS['备注'], SALE_SHOW['校验'], SALE_LAST, N_SALE,
     '客户不在往来单位、产品不在物料档案、吨位金额不是数字、有运费没物流单位；⚠ 还没定价（收入先不算）'),
    (SH_PROD, '产', PROD_COLS['备注'], PROD_SHOW['校验'], PROD_LAST, N_PROD,
     '产品不在物料档案、原料列在物料档案里没指定是哪种料、数字写成了文字、日期不在本年'),
    (SH_WAGE, '工', WAGE_COLS['备注'], WAGE_SHOW['校验'], WAGE_LAST, N_WAGE,
     '月份看不懂或不在本年、姓名不在【基础资料】⑤人员、部门不对、应发是 0'),
    (SH_INV, '票', INV_COLS['备注'], INV_SHOW['校验'], INV_LAST, N_INV,
     '没填往来单位或价税合计、方向不对、票种不在⑥票种；⚠ 垫付编号找不到、发票号码重复'),
    (SH_CNT, '盘', CNT_COLS['备注'], CNT_SHOW['校验'], CNT_LAST, N_CNT,
     '物料不在档案里、月份看不懂、实盘数量不是数字；同一种料同一个月盘了两次只认第一条'),
    (SH_MJ, '手', MJ_COLS['备注'], MJ_SHOW['校验'], MJ_LAST, N_MJ,
     '科目不在会计科目里、同一组借贷不平或日期不一样、记了现金银行或存货科目（不准）；整组平了才记账'),
]
INPUT_COLS = {SH_CASH: CASH_COLS, SH_PUR: PUR_COLS, SH_SALE: SALE_COLS, SH_PROD: PROD_COLS, SH_WAGE: WAGE_COLS,
              SH_INV: INV_COLS, SH_CNT: CNT_COLS, SH_MJ: MJ_COLS}


def last_input(sh):
    return max(INPUT_COLS[sh].values(), key=CI)


def chk_end(cap):
    """「这一行的问题」、备注 列的计数区域：表头行到容量后 2000 行（跟 s_input 一样）"""
    return R0 + cap + 2000


# 隐藏辅助列（从 M 起，一个键一列；第 i 个在第 i+1 行）
N_CHUNKS = {SRC_CASH: math.ceil(N_CASH / CHUNK), SRC_PUR: math.ceil(N_PUR / CHUNK), SRC_SALE: math.ceil(N_SALE / CHUNK)}
N_AB = sum(N_CHUNKS.values()) + 1         # ＋ 手工整块
HELPERS = [('垫备期初', N_ACC), ('人员重名', N_PER), ('账户重名', N_ACC), ('类别重名', N_CAT), ('报销重名', N_RE),
           ('票种重名', N_TK), ('借期无效', N_LO), ('料期初', N_RM), ('品期初', N_PRD), ('品料重名', N_PRD),
           ('矩月差', 12), ('月末凭证差', len(ME_TYPES) * 12), ('段差', N_AB), ('上级记账', N_COA), ('利息多付', N_LEND),
           ('凭证前', 13), ('凭证张', 12), ('负料', N_RM), ('负品', N_PRD), ('工资没登', 12), ('账户负', N_ACC),
           ('已补票', N_INV), ('卡折旧', 12), ('预付没摊', N_OP)]
HCOL = {k: CL(CI('M') + i) for i, (k, _n) in enumerate(HELPERS)}
HN = dict(HELPERS)


def hrng(k):
    c = HCOL[k]
    return f'${c}$2:${c}${HN[k] + 1}'


def seg_list():
    """凭证段：[(说明, _分 第一行, 最后一行)]：月末凭证逐张（代码×月），资金/进货/销售每 CHUNK 张一段，手工整块"""
    me, ab = [], []
    for code, nm, legs in ME_TYPES:
        b = FEN_BASE[code]
        for m in range(1, 13):
            a = b + (m - 1) * legs
            me.append((f'{m}月{nm}', a, a + legs - 1))
    for src in (SRC_CASH, SRC_PUR, SRC_SALE):
        b, legs = FEN_BASE[src], FEN_LEGS[src]
        cnt = {SRC_CASH: N_CASH, SRC_PUR: N_PUR, SRC_SALE: N_SALE}[src]
        for c in range(N_CHUNKS[src]):
            n1, n2 = c * CHUNK + 1, min((c + 1) * CHUNK, cnt)
            ab.append((f'{src}:{n1}-{n2}', b + (n1 - 1) * legs, b + n2 * legs - 1))
    b = FEN_BASE[SRC_MJ]
    ab.append(('手工', b, b + N_MJ - 1))
    return me, ab


def bal(code):
    """截止日那个月月末的科目余额（借正贷负，一级＝含下级）：矩_汇期初＋Σ1～截止月(汇借－汇贷)"""
    k = kh(code)
    return (f'IF({k}=0,0,ROUND(INDEX(矩_汇期初,{k})+SUM(INDEX(矩_汇借,{k},1):INDEX(矩_汇借,{k},P_截止月))'
            f'-SUM(INDEX(矩_汇贷,{k},1):INDEX(矩_汇贷,{k},P_截止月)),2))')


def over_rows(sh, hdr_row, cap, c2):
    """表头行往下第 cap+1 行起（A～c2 列）还有几格有东西"""
    return (f'COUNTA(INDEX({sh}!$A:$A,ROW({sh}!$A${hdr_row})+{cap + 1}):INDEX({sh}!${c2}:${c2},1048576))')


def over_base():
    """基础资料每块：第 cap 个以后、下一块标题以前（最后一块：往下 300 行）A 列还有没有名字"""
    parts = []
    for key, _t, _tr, hr, cap in BASE_BLOCKS:
        a = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+{cap + 1})'
        b = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+P_{key}位)'
        parts.append(f'IF(P_{key}位<={cap},0,COUNTA({a}:{b}))')
    return '+'.join(parts)


def over_item():
    a = (f'IF(ROW({SH_ITEM}!$A${ITEM_PRD_R0})-1>=ROW({SH_ITEM}!$A${ITEM_RM_HDR})+{N_RM + 1},'
         f'COUNTA(INDEX({SH_ITEM}!$A:$A,ROW({SH_ITEM}!$A${ITEM_RM_HDR})+{N_RM + 1}):'
         f'INDEX({SH_ITEM}!${RM_COLS["备注"]}:${RM_COLS["备注"]},ROW({SH_ITEM}!$A${ITEM_PRD_R0})-1)),0)')
    b = over_rows(SH_ITEM, ITEM_PRD_HDR, N_PRD, PRD_COLS['备注'])
    return f'{a}+{b}'


BASE_SHEETS = [(SH_BASE, 'M'), (SH_UNIT, UNIT_LAST), (SH_ITEM, ITEM_SHOW_RM['校验']), (SH_FA, FA_LAST), (SH_FIX, FIX_LAST),
               (SH_COA, COA_LAST)]


def cols_bad(sh, last):
    return f'IF(COLUMNS({sh}!$A${HDR}:${last}${HDR})<>{CI(last)},1,0)'


def base_ex():
    """基础表里备注/说明写「示例」开头的行"""
    blk = lambda key, col, cap: (f'COUNTIF(INDEX({inref(SH_BASE, col, block=key)},2):INDEX({inref(SH_BASE, col, block=key)},'
                                 f'{cap + 1}),"示例*")')
    return '+'.join([
        f'COUNTIF({SH_UNIT}!${UNIT_COLS["备注"]}${HDR}:${UNIT_COLS["备注"]}${R0 + N_UNIT + 200},"示例*")',
        blk('人员', PER_COLS['备注'], N_PER), blk('借期', LO_COLS['说明'], N_LO), blk('期初', OP_COLS['说明'], N_OP),
        f'COUNTIF({SH_FA}!${FA_COLS["备注"]}${HDR}:${FA_COLS["备注"]}${R0 + N_FA + 200},"示例*")',
        f'COUNTIF({SH_FIX}!${FIX_COLS["备注"]}${HDR}:${FIX_COLS["备注"]}${R0 + N_FIX + 200},"示例*")'])


def raw_nonzero(x):
    """录入格填了东西且不是 0"""
    return f'IF(ISNUMBER({x}),IF({x}<>0,1,0),IF(TRIM({x}&"")<>"",1,0))'


# ───────────────────────── 辅助列公式（第 i 个，1 起） ─────────────────────────
def helper_formulas():
    acc_raw = lambda i: f'INDEX({inref(SH_BASE, ACC_COLS["期初余额"], block="账户")},{i}+1)'
    dup = lambda rng, i: f'IF(INDEX({rng},{i})="",0,IF(COUNTIF({rng},{esc(f"INDEX({rng},{i})")})>1,1,0))'
    lo = lambda k, i: f'INDEX({inref(SH_BASE, LO_COLS[k], block="借期")},{i}+1)'
    rmx = lambda k, i: f'INDEX({inref(SH_ITEM, RM_COLS[k], block="原料")},{i}+1)'
    pdx = lambda k, i: f'INDEX({inref(SH_ITEM, PRD_COLS[k], block="产品")},{i}+1)'
    me, ab = seg_list()
    fe, ff = FEN_COLS['借'], FEN_COLS['贷']
    seg = lambda a, b: f'ROUND(SUM({H_FEN}!${fe}${a}:${fe}${b})-SUM({H_FEN}!${ff}${a}:${ff}${b}),2)'
    xi = lambda pre, i: (f'ROUND(INDEX(息_{pre}期初应付息,{i})+INDEX(息_{pre}累计利息,{i})-INDEX(息_{pre}已付息,{i}),2)<-0.005')
    neg = lambda nm, ok, i: (f'IF(INDEX({ok},{i})<>1,0,IF(COUNTIF(INDEX({nm},{i},1):INDEX({nm},{i},P_截止月),"<-0.0001")>0,1,0))')
    k2211 = kh('2211')
    # 第 m 月：付了应付职工薪酬（资金台帐对方科目 2211 净流出）、工资表这个月没有有效行、上个月也没有（建账月：⑧ 也没有期初应付工资）
    wage = lambda m: (f'IF(AND({m}>=P_建账月,{m}<=P_截止月),IF(AND(SUMIFS(资_净额,资_原对方号,{k2211},资_月,{m},资_有效,1)<0,'
                      f'COUNTIFS(工_月,{m},工_有效,1)=0,NOT(IF({m}=P_建账月,INDEX(科目_期初,MAX(1,{k2211}))<0,'
                      f'COUNTIFS(工_月,{m - 1},工_有效,1)>0))),1,0),0)')
    zr = lambda i: f'INDEX(票_补票行,{i})'
    paid = lambda i: (f'IF(OR(INDEX(票_有效,{i})<>1,{zr(i)}=0),0,IF(COUNTIFS(INDEX(票_补票行,1):{zr(i)},{zr(i)},'
                      f'INDEX(票_有效,1):INDEX(票_有效,{i}),1)>1,0,IF(INDEX(资_日期,{zr(i)})>{D},0,-INDEX(资_净额,{zr(i)}))))')
    return {
        '垫备期初': lambda i: (f'IF(OR(INDEX(账户_类型号,{i})=5,INDEX(账户_类型号,{i})=6),{raw_nonzero(acc_raw(i))},0)'),
        '人员重名': lambda i: dup('人员_姓名', i),
        '账户重名': lambda i: dup('账户_名称', i),
        '类别重名': lambda i: dup('类别_名称', i),
        '报销重名': lambda i: dup('报销_名称', i),
        '票种重名': lambda i: dup('票种_名称', i),
        '借期无效': lambda i: (f'IF({i}>P_借期位,0,IF(INDEX(借期_有效,{i})=1,0,IF(OR('
                              + ','.join(f'TRIM({lo(k, i)}&"")<>""' for k in ('出借人', '日期', '金额', '性质', '已付利息'))
                              + '),1,0)))'),
        '料期初': lambda i: (f'IF(INDEX(料档_期初量,{i})=0,0,IF(OR(ISNUMBER({rmx("期初单价", i)}),ISNUMBER({rmx("期初金额", i)})),0,1))'),
        '品期初': lambda i: (f'IF(INDEX(品档_期初量,{i})=0,0,IF(OR(ISNUMBER({pdx("期初单价", i)}),ISNUMBER({pdx("期初金额", i)})),0,1))'),
        '品料重名': lambda i: (f'IF(INDEX(品档_有效,{i})<>1,0,IF(COUNTIF(料档_名称,{esc(f"INDEX(品档_名称,{i})")})>0,1,0))'),
        '矩月差': lambda m: f'ROUND(SUM(INDEX(矩_借,0,{m}))-SUM(INDEX(矩_贷,0,{m})),2)',
        '月末凭证差': lambda i: seg(me[i - 1][1], me[i - 1][2]),
        '段差': lambda i: seg(ab[i - 1][1], ab[i - 1][2]),
        '上级记账': lambda r: (f'IF(AND(INDEX(科目_有效,{r})=1,INDEX(科目_级次,{r})=1),IF(COUNTIFS(科目_一级号,{r},科目_级次,2,'
                              f'科目_有效,1)=0,0,IF(OR(INDEX(科目_期初,{r})<>0,SUM(INDEX(矩_借,{r},0))<>0,SUM(INDEX(矩_贷,{r},0))<>0),1,0)),0)'),
        '利息多付': lambda l: f'IF(INDEX(出借_单位号,{l})=0,0,IF(OR({xi("借", l)},{xi("投", l)}),1,0))',
        '凭证前': lambda m: f'COUNTIF(凭_键,"<"&DATE(P_年度,{m},1)*{KEY_M})',
        '凭证张': lambda m: f'{HCOL["凭证前"]}{m + 2}-{HCOL["凭证前"]}{m + 1}',
        '负料': lambda k: neg('料_末量', '料档_有效', k),
        '负品': lambda p: neg('品_末量', '品档_有效', p),
        '工资没登': wage,
        '账户负': lambda i: (f'IF(AND(INDEX(账户_类型号,{i})>=1,INDEX(账户_类型号,{i})<=4),IF(ROUND(INDEX(账户_期初,{i})'
                            f'+SUMIFS(资_净额,资_账户号,{i},资_日期,"<="&{D}),2)<0,1,0),0)'),
        '已补票': paid,
        '卡折旧': lambda m: f'IF({m}>P_截止月,0,SUMIFS(INDEX(卡_月折旧块,0,{m}),卡_有效,1,卡_处置月,">"&P_截止月))',
        # ⑧ 1123（预付）有余额，但摊进没认出来或剩余月数空：不摊（当普通余额）
        '预付没摊': lambda i: (f'IF(AND(INDEX(期初_有效,{i})=1,INDEX(期初_科目号,{i})>0),IF(AND(INDEX(科目_一级号,'
                              f'MAX(1,INDEX(期初_科目号,{i})))={kh("1123")},ROUND(INDEX(期初_借,{i})-INDEX(期初_贷,{i}),2)<>0,'
                              f'OR(INDEX(期初_摊进号,{i})=0,INDEX(期初_剩余月数,{i})=0)),1,0),0)'),
    }


# ───────────────────────── 检查项 ─────────────────────────
def first_name(hk, names):
    return f'IFERROR(INDEX({names},MATCH(1,{hrng(hk)},0)),"")'


def ye_cells(wb):
    """年末结转里各块下面的核对行（s_je 生成的）：A 列公式里有「跟账上一致」或「⑧ 一共」"""
    ws = wb[SH_YE]
    out = []
    for row in ws.iter_rows(min_col=1, max_col=1):
        c = row[0]
        if isinstance(c.value, str) and c.value.startswith('=') and ('跟账上一致' in c.value or '⑧ 一共' in c.value):
            out.append(c.coordinate)
    return out


def items(wb):
    """②③④ 每一项：dict(key, sec, label, F, G, H, kind, cond, link, note, fmt, text)；text＝F:H 合并放一段文字（引用别的表的核对格）"""
    has = lambda nm: nm in wb.defined_names
    raw_d0 = f'{SH_BASE}!${BASE_PAR_VAL}${PAR_ROW["建账日"]}'
    k_fix = f'{H_KM}!${KM_FIXED["号"]}$2:${KM_FIXED["号"]}${len(FIXED_CODES) + 1}'
    k_fixc = f'{H_KM}!${KM_FIXED["代码"]}$2:${KM_FIXED["代码"]}${len(FIXED_CODES) + 1}'
    k4101 = kh('4101')
    n_oh = f'COUNTIFS(科目_一级号,{k4101},科目_级次,2,科目_有效,1)'
    n_fee = '+'.join(f'COUNTIFS(科目_一级号,{kh(p)},科目_级次,2,科目_有效,1)' for p in FEE_PARENTS)
    cols_names = '&'.join(f'IF({cols_bad(sh, last)}=1,"、{sh}","")' for sh, last in BASE_SHEETS)
    k1122, k2202 = kh('1122'), kh('2202')
    def late(nm, dnm, *cond):
        """截止日以后、截止月月末以前（_位 只算到截止日；科目余额算到月末）"""
        cs = ''.join(c + ',' for c in cond)
        return f'SUMIFS({nm},{cs}{dnm},">"&{D},{dnm},"<="&{ME})'

    pos = lambda nm: nm + ',">0"'
    mj_late = lambda side, k: late(f'手_{side}', '手_日期', pos('手_单位号'), f'手_一级号,{k}', '手_记账,1')
    ar_sum = ('ROUND(SUM(位_应收)+' + late('销_收入含税', '销_日期', pos('销_客户号'))
              + '+' + late('资_应收变动', '资_日期', pos('资_单位号')) + '+' + mj_late('借', k1122) + '-' + mj_late('贷', k1122) + ',2)')
    ap_sum = ('ROUND(SUM(位_应付)+' + late('进_货款计', '进_日期', pos('进_供应商号')) + '+' + late('进_运费计', '进_日期', pos('进_物流号'))
              + '+' + late('销_运费计', '销_日期', pos('销_物流号')) + '+' + late('资_应付变动', '资_日期', pos('资_单位号'))
              + '+' + mj_late('贷', k2202) + '-' + mj_late('借', k2202) + ',2)')
    eng = lambda nm: f'ROUND(SUM(INDEX({nm},0,P_截止月)),2)'
    card_cost = f'ROUND(SUMIFS(卡_原值,卡_有效,1,卡_取得日,"<="&{ME},卡_处置月,">"&P_截止月),2)'
    card_dep = (f'ROUND(SUMIFS(卡_建账前已提,卡_有效,1,卡_取得日,"<"&P_建账日,卡_处置月,">"&P_截止月)+SUM({hrng("卡折旧")}),2)')
    fix = '(固_本月生效=1)*(固_每月几号>0)*(固_每月几号<DAY(P_截止))*(固_本月已付<固_每月金额-0.005)'
    ye = ye_cells(wb) if SH_YE in wb.sheetnames else []
    ye_n = '+'.join(f'IF(OR(ISNUMBER(FIND("✗",{SH_YE}!{a})),ISNUMBER(FIND("⚠",{SH_YE}!{a}))),1,0)' for a in ye) or '0'
    ref = lambda nm: nm if has(nm) else None
    L = []

    def it(sec, key, label, H=None, F=None, G=None, kind='✗', cond='{H}>0', link=None, note='', fmt='0', fmtF=None,
           fmtG=None, text=None):
        L.append(dict(sec=sec, key=key, label=label, H=H, F=F, G=G, kind=kind, cond=cond, link=link, note=note, fmt=fmt,
                      fmtF=fmtF or fmt, fmtG=fmtG or fmt, text=text))

    # ── ② 基础资料 ──
    S = 2
    it(S, '建账日', '建账日期没填、看不懂或不是某月 1 号',
       H=(f'IF(OR(TRIM({raw_d0}&"")="",AND(ISNUMBER({raw_d0}),{raw_d0}<36526,OR({raw_d0}<200001,{raw_d0}>209912)),'
          f'AND(NOT(ISNUMBER({raw_d0})),ISERROR(DATEVALUE(TRIM({raw_d0}&"")))),DAY(P_建账日)<>1),1,0)'),
       F='P_建账日', fmtF=DATE, link=SH_BASE,
       note='【基础资料】① 参数「建账日期」填某月 1 号（比如 2026-08-01），期初余额都是这天前一晚的数；不是 1 号，第一个月的折旧、利息、费用都不对。F 列是现在认出来的日期')
    it(S, '垫备期初', '② 账户：垫付、备用金账户填了期初余额（不算）', H=f'SUM({hrng("垫备期初")})',
       F=first_name('垫备期初', '账户_名称'), fmtF='@', link=SH_BASE,
       note='垫付（老板垫付）、备用金账户的期初余额不填：建账前老板垫的钱填 ⑦（性质选垫付），建账前给员工的备用金填 ⑧（科目 122103，明细填人名）')
    it(S, '借期无效', '⑦ 建账前的借款、投资、垫付：有内容但不算的行', H=f'SUM({hrng("借期无效")})',
       F=f'IFERROR("第 "&MATCH(1,{hrng("借期无效")},0)&" 行","")', fmtF='@', link=SH_BASE,
       note='⑦ 每一行要：出借人是【往来单位】里类型＝出借人的、日期不晚于建账日、金额不是 0、性质选借款/投资/垫付；不然这一行不算（期初借款、应付利息都会少）')
    it(S, '期初无效', '⑧ 其他期初余额：填了金额但不算的行',
       H='COUNTIFS(期初_有效,0,期初_借,"<>0")+COUNTIFS(期初_有效,0,期初_借,0,期初_贷,"<>0")', link=SH_BASE,
       note='科目编码要在【会计科目】里；⑧ 不能填：现金银行微信票据（填 ② 账户期初）、应收应付（往来单位期初）、原材料库存商品（物料档案期初）、'
            '生产成本（产品「期初在产」）、固定资产累计折旧（卡片）、借款实收资本垫付应付利息（⑦）、未分配利润（自动倒挤）、损益类科目')
    it(S, '期初预付', '⑧ 预付账款（1123）：没按月摊的行（摊进、剩余月数没填或认不出来，当普通余额）',
       H=f'SUM({hrng("预付没摊")})', F=f'IFERROR("第 "&MATCH(1,{hrng("预付没摊")},0)&" 行","")', fmtF='@', kind='⚠',
       link=SH_BASE,
       note='一次付几个月的预付（房租等）要在新本接着按月摊：⑧ 这一行「剩余摊销月数」填还要摊几个月，「摊进」填摊到哪个费用明细科目'
            '（编码或名称，如 410108 或 厂房租金；要是制造费用、管理费用、销售费用、财务费用、主营业务成本下面的明细）。'
            '不用摊的预付（比如预付货款）可以不管。F 列是第一处在 ⑧ 第几行')
    it(S, '期初平', '期初试算平衡（全部科目期初 借方－贷方，差多少）', H='ROUND(SUM(科目_期初),2)', cond='ABS({H})>=0.01',
       fmt=MONEY, link=SH_TB,
       note='未分配利润是倒挤的，正常一定平；不平说明【会计科目】里少了固定科目（看下面「少了固定科目」），或者改了科目编码')
    it(S, '倒挤', '期初未分配利润（倒挤出来的，贷方为正）', H=f'IF({kh("310415")}=0,0,-INDEX(科目_期初,{kh("310415")}))',
       kind='ℹ', cond='{H}<>0', fmt=MONEY, link=SH_BS,
       note='＝期初资产－负债－实收资本；数不对多半是期初少填了哪一项（② 账户余额、往来单位期初、物料期初、固定资产卡片、⑦、⑧）')
    it(S, '物料期初', '物料档案：期初有数量没单价、没金额（期初单价明写 0 的不算）',
       H=f'SUM({hrng("料期初")})+SUM({hrng("品期初")})',
       F=f'IFERROR(INDEX(料档_名称,MATCH(1,{hrng("料期初")},0)),{first_name("品期初", "品档_名称")})', fmtF='@', link=SH_ITEM,
       note='期初数量一定要有期初单价或期初金额，不然期初库存是 0 元，成本、利润都算错；真的不值钱（比如废料）期初单价填 0')
    it(S, '单位重名', '往来单位：简称重复', H='COUNTIF(位_校验,"✗ 简称重复*")', link=SH_UNIT,
       note='同一家只建一行（别的叫法写在「别名」里）；重复的那一行不算，录入表选它会对不上')
    it(S, '物料重名', '物料档案：原料或产品名称重复（含产品跟原料重名）',
       H=f'COUNTIFS(料档_名称,"?*",料档_有效,0)+COUNTIFS(品档_名称,"?*",品档_有效,0)+SUM({hrng("品料重名")})', link=SH_ITEM,
       note='一种料一行、名字不能重复；产品名不要跟原料一样（销售、进货按名字认料还是认产品）')
    it(S, '人员重名', '【基础资料】⑤ 人员：姓名重复', H=f'SUM({hrng("人员重名")})', F=first_name('人员重名', '人员_姓名'),
       fmtF='@', link=SH_BASE, note='重名的加个字区分（报销人、工资表按姓名认人）')
    it(S, '清单重名', '【基础资料】② 账户、③ 收支类别、④ 报销项目、⑥ 票种：名称重复',
       H=f'SUM({hrng("账户重名")})+SUM({hrng("类别重名")})+SUM({hrng("报销重名")})+SUM({hrng("票种重名")})', link=SH_BASE,
       note='同一个名字只写一行（只认上面那一行，下面那行的设置不起作用）')
    it(S, '固定科目', '【会计科目】少了固定科目（编码被改了或删了）', H=f'COUNTIF({k_fix},0)',
       F=f'IFERROR(INDEX({k_fixc},MATCH(0,{k_fix},0)),"")', fmtF='@', link=SH_COA,
       note='F 列是第一个找不到的编码。预置科目的编码不许改、不许删（可以改名、加二级）；改回来就好')
    it(S, '科目坏', '【会计科目】编码有问题的行（重复、没有上级、不是 4 位/6 位数字）', H='COUNTIFS(科目_编码,"?*",科目_有效,0)',
       kind='⚠', link=SH_COA, note='这些科目下拉里没有、记不了账；看【会计科目】最右边「这一行的问题」')
    it(S, '单位错', '往来单位：别的 ✗ 行（没填类型、全称重复、出借人填了期初应收应付…）',
       H='COUNTIF(位_校验,"✗*")-COUNTIF(位_校验,"✗ 简称重复*")', link=SH_UNIT, note='看【往来单位】最右边「这一行的问题」')
    it(S, '卡片错', '固定资产：有 ✗ 的卡片', H='COUNTIF(卡_校验,"✗*")', link=SH_FA, note='看【固定资产】最右边「这一行的问题」（日期、原值、月数）')
    it(S, '固定错', '固定支出：有 ✗ 的行', H='COUNTIF(固_校验,"✗*")', link=SH_FIX, note='看【固定支出】最右边「这一行的问题」')
    it(S, '基础警', '往来单位、固定资产、固定支出：有 ⚠ 的行', kind='⚠',
       H='COUNTIF(位_校验,"⚠*")+COUNTIF(卡_校验,"⚠*")+COUNTIF(固_校验,"⚠*")', link=SH_UNIT,
       note='看各表最右边「这一行的问题」（预付了钱的供应商、出借人类别没选…）')
    it(S, '基础超', '基础表超容量（多出来的不算）',
       H='+'.join([over_base(), over_rows(SH_UNIT, HDR, N_UNIT, UNIT_COLS['备注']), over_item(),
                   over_rows(SH_FA, HDR, N_FA, FA_COLS['备注']), over_rows(SH_FIX, HDR, N_FIX, FIX_COLS['备注']),
                   over_rows(SH_COA, HDR, N_COA, COA_COLS['说明'])]), link=SH_BASE,
       note=f'账户最多 {N_ACC} 个、收支类别 {N_CAT}、报销项目 {N_RE}、人员 {N_PER}、票种 {N_TK}、⑦ {N_LO} 行、⑧ {N_OP} 行、往来单位 {N_UNIT} 家、'
            f'原料 {N_RM} 种、产品 {N_PRD} 种、固定资产 {N_FA} 张、固定支出 {N_FIX} 条、会计科目 {N_COA} 个；数是格子数')
    it(S, '基础列', '基础表插了列或删了列', H='+'.join(cols_bad(sh, last) for sh, last in BASE_SHEETS),
       F=f'MID({cols_names},2,200)', fmtF='@', link=SH_BASE,
       note='基础资料、往来单位、物料档案、固定资产、固定支出、会计科目的列不能插、不能删（公式按列取数）；撤销或把列删掉/补回来')
    it(S, '基础示例', '基础表里的示例行（往来单位、人员、⑦、⑧、固定资产、固定支出）', H=base_ex(), kind='ℹ', link=SH_UNIT,
       note='灰底、备注写「示例」开头的行：正式用之前整行删掉（往来单位删了，录入表里的示例行也要一起删）')

    # ── ③ 账 ──
    S = 3
    it(S, '矩阵平', '科目×月汇总：每个月借方合计＝贷方合计（不平的月数）', H=f'COUNTIF({hrng("矩月差")},"<>0")', link=SH_TB,
       note='不平多半是有分录记到了科目表里没有的科目（看【科目余额表】F3）；正常一定平，不平请找我')
    it(S, '凭证平', '每张凭证借＝贷（不平的张数；资金、进货、销售按每 50 张一段查）',
       H=f'COUNTIF({hrng("月末凭证差")},"<>0")+COUNTIF({hrng("段差")},"<>0")', link=SH_VLIST,
       note='每张凭证都是按公式拆的分录，正常一定平；不平请找我（【凭证汇总】借贷平衡一列能看到是哪张）')
    if SH_TB in wb.sheetnames:
        it(S, '余额表', '科目余额表平衡检查（F3，B3 选的月份）', text=f'{SH_TB}!$F$3', cond='LEFT({T}&"",1)<>"√"', link=SH_TB,
           note='月初、本月、本年、月末借贷都要平，还查有没有分录记到无效科目、期初平不平')
    for nm, lab, lk, nt in (
            ('资产负债表_平衡', '="资产负债表：期末、年初 资产＝负债＋所有者权益（"&资产负债表_选月&" 月）"', SH_BS,
             '资产负债表 D3 选的月份（空着＝截止日那个月）；不平先看【科目余额表】平不平'),
            ('资产负债表_往来核对', '="资产负债表：客户、供应商逐家余额合计＝应收账款、应付账款（"&资产负债表_选月&" 月）"', SH_BS,
             '对不上多半是资金台帐、手工凭证记了应收应付但往来单位没填或不在【往来单位】里'),
            ('利润表_核对', '利润表：净利润＝损益类科目 贷方－借方', SH_PL, '对不上多半是自己加的损益科目编码不在报表行里'),
            ('现金流量表_核对', '现金流量表：现金流量合计＝货币资金变动；期末现金＝资产负债表货币资金', SH_CF,
             '对不上多半是收支类别的「现金流量」填的不是 1～18 的行次')):
        if ref(nm):
            it(S, nm, lab, text=nm, cond='LEFT({T}&"",1)<>"√"', link=lk, note=nt)
    it(S, '应收', '应收账款 1122：往来单位逐家合计＝科目余额（截止日那个月月末）', F=ar_sum, G=bal('1122'), H='ROUND({F}-{G},2)',
       cond='ABS({H})>=0.01', fmt=MONEY, link=SH_AR,
       note='F＝逐家合计（客户往来）、G＝科目余额；差的是资金台帐、手工凭证里记了应收但往来单位没填或不在【往来单位】的（那几行本身也是 ✗）')
    it(S, '应付', '应付账款 2202：往来单位逐家合计＝科目余额（贷方为正）', F=ap_sum, G=f'-{bal("2202")}', H='ROUND({F}-{G},2)',
       cond='ABS({H})>=0.01', fmt=MONEY, link=SH_AP,
       note='F＝逐家合计（供应商往来）、G＝科目余额；差的是付货款没选往来单位、手工凭证 2202 没填明细的')
    for key, code, nm, lab, lk, nt in (
            ('1403', '1403', '料_末额', '原材料 1403 ＝ 进销存原料（含回收料）月末金额合计', SH_STK, '进货、领料、废料回收、卖料、盘点都按进销存记账，正常一定相等；不等请找我'),
            ('1405', '1405', '品_末额', '库存商品 1405 ＝ 进销存产品月末金额合计', SH_STK, '完工入库、外购、销售成本、盘点，正常一定相等；不等请找我'),
            ('4001', '4001', '本_末在产', '生产成本 4001 ＝ 各产品月末在产（投了料还没产出的）', SH_COST, '正常一定相等；不等请找我'),
            ('4101', '4101', '本池_留下月', '制造费用 4101 ＝ 各明细池留到下月的（当月一吨没产就整月留着）', SH_COST,
             '不等多半是收支类别、手工凭证直接记在一级科目 4101 上（看下面「上级科目直接记账」），或者 ⑧ 填了 4101')):
        it(S, key, lab, F=bal(code), G=eng(nm), H='ROUND({F}-{G},2)', cond='ABS({H})>=0.01', fmt=MONEY, link=lk,
           note='F＝科目余额、G＝' + {'1403': '进销存', '1405': '进销存', '4001': '在产', '4101': '池留存'}[key] + '；' + nt)
    it(S, '1601', '固定资产 1601 ＝ 卡片原值合计（没处置的）', F=bal('1601'), G=card_cost, H='ROUND({F}-{G},2)', cond='ABS({H})>=0.01',
       kind='⚠', fmt=MONEY, link=SH_FA,
       note='F＝科目余额、G＝卡片；买设备付了钱要到【固定资产】登卡片（原值填不含税）；只付了定金的先不用管；处置了要做手工凭证冲 1601、1602')
    it(S, '1602', '累计折旧 1602 ＝ 卡片累计折旧合计（贷方为正）', F=f'-{bal("1602")}', G=card_dep, H='ROUND({F}-{G},2)',
       cond='ABS({H})>=0.01', kind='⚠', fmt=MONEY, link=SH_FA,
       note='F＝科目余额、G＝卡片（建账前已提＋本年各月折旧）；处置的卡片要做手工凭证冲掉；建账以后买的卡片「建账前已提」不要填')
    it(S, '122104', '内部转账待对冲 122104 ＝ 0', H=bal('122104'), cond='ABS({H})>=0.01', fmt=MONEY, link=SH_CASH,
       note='自己账户之间倒钱要记两行（转出一行支出、转进一行收入，金额一样）；资金台帐按类别筛「内部转账」对一下')
    it(S, '122105', '承兑托收贴现待对冲 122105 ＝ 0', H=bal('122105'), cond='ABS({H})>=0.01', fmt=MONEY, link=SH_CASH,
       note='承兑到期、贴现要记两行：承兑汇票账户一行支出、银行一行收入；贴现少的钱在承兑账户记「贴现息」')
    it(S, '利息多付', '应付利息 2231 / 应付利润 2232：利息付多了的出借人（借方余额）', H=f'SUM({hrng("利息多付")})',
       F=first_name('利息多付', '出借_名称'), fmtF='@', kind='⚠', link=SH_LOAN,
       note='付的利息比计提的多：利率、计息天数、⑦「建账前已付利息」对不对？先付后算的到月底计提了就好')
    it(S, '上级记账', '上级科目直接记账（有二级科目的一级科目有发生额或期初）', H=f'SUM({hrng("上级记账")})',
       F=f'IFERROR(INDEX(科目_编码,MATCH(1,{hrng("上级记账")},0))&" "&INDEX(科目_名称,MATCH(1,{hrng("上级记账")},0)),"")',
       fmtF='@', kind='⚠', link=SH_TB,
       note='F 是第一个；收支类别、报销项目、手工凭证、⑧ 的科目选到二级（明细）：费用分析、制造费用按明细分，记在一级上分不进去')
    it(S, '凭证张', f'凭证张数：一个月超过 {N_VLIST} 张的月数（F＝最多的那个月几张）', F=f'MAX({hrng("凭证张")})',
       H=f'COUNTIF({hrng("凭证张")},">{N_VLIST}")', link=SH_VLIST,
       note=f'【凭证汇总】一个月最多列 {N_VLIST} 张，多的打印不出来：找我加容量')
    if SH_VLIST in wb.sheetnames:
        it(S, '汇总K3', '凭证汇总（K3，B3 选的月份）：张数、借贷不平、打印页数',
           text=(f'IF({SH_VLIST}!$K$3&""="","√ "&{"汇总_选月" if has("汇总_选月") else "P_截止月"}&" 月：张数、借贷都没问题",'
                 f'{SH_VLIST}!$K$3&"")'),
           cond='LEFT({T}&"",1)<>"√"', kind='IF(LEFT({T}&"",1)="✗","✗","⚠")', link=SH_VLIST,
           note=f'超过 {N_VPAGE} 页的：【记账凭证】K2「从第几页起」分批印')
    it(S, '制费槽', f'制造费用 4101 的二级明细超过 {N_OH} 个', F=n_oh, H=f'MAX(0,{n_oh}-{N_OH})', link=SH_COA,
       note=f'制造费用按明细各自成池分配，最多 {N_OH} 个明细，多的分不出去：不用的明细删掉或合并')
    it(S, '费用槽', f'费用明细（4101/5401/5601/5602/5603 的二级）超过 {N_FEE} 个', F=n_fee, H=f'MAX(0,{n_fee}-{N_FEE})',
       link=SH_COA, note=f'摊销、跨月费用归属、补票冲减按费用明细记账，最多 {N_FEE} 个，多的记不进去：不用的明细删掉或合并')
    if ye:
        it(S, '年末结转', '年末结转：各块跟账上核对（对不上的块数）', H=ye_n,
           kind='IF(P_截止月<12,"ℹ","⚠")', link=SH_YE,
           note='到 12 月底才用；截止日期不在 12 月时只是提示（月末计提没做完）。处置了卡片没做凭证，固定资产那块会对不上')

    # ── ④ 业务提醒 ──
    S = 4
    it(S, '负库存', '负库存：原料、产品哪个月末数量是负数（种数）', H=f'SUM({hrng("负料")})+SUM({hrng("负品")})',
       F=f'IFERROR(INDEX(料档_名称,MATCH(1,{hrng("负料")},0)),{first_name("负品", "品档_名称")})', fmtF='@', link=SH_STK,
       note='F 是第一种。多半是进货没登、期初没填、生产日报领错了列或多写了；领超的照单价出库，成本会不准')
    it(S, '工资没登', '发了工资（应付职工薪酬）这个月和上个月都没登工资表（月数）', H=f'SUM({hrng("工资没登")})',
       F=f'IFERROR(MATCH(1,{hrng("工资没登")},0)&" 月","")', fmtF='@', link=SH_WAGE,
       note='资金台帐记了「员工工费」「员工借支」，工资表就要登那个月（或上个月）的工资，不然工资成本少算、应付职工薪酬变成借方')
    it(S, '待定价进', '待定价的进货（笔；F＝按参考单价暂估的金额）',
       H=f'COUNTIFS(进_待定价,1,进_有效,1,进_日期,"<="&{D})', F=f'SUMIFS(进_暂估额,进_待定价,1,进_有效,1,进_日期,"<="&{D})',
       kind='⚠', fmtF=MONEY, link=SH_PUR, note='进货物流登记里单价、货款空着的：成本先按参考单价暂估；定了价回原行补上（日期别改）')
    it(S, '待定价销', '待定价的发货（笔；F＝吨位）', H=f'COUNTIFS(销_待定价,1,销_有效,1,销_日期,"<="&{D})',
       F=f'SUMIFS(销_吨位,销_待定价,1,销_有效,1,销_日期,"<="&{D})', kind='⚠', fmtF='#,##0.###', link=SH_SALE,
       note='销售台账里单价、金额空着的：出库和成本照算，收入、应收还没算；定了价回原行补上')
    it(S, '待补票', '待补票的付款（笔；F＝金额；垫付、备用金这些先付钱后补票的）',
       H=f'COUNTIFS(资_待补票,1,资_日期,"<="&{D})-COUNTIF({hrng("已补票")},"<>0")',
       F=f'ROUND(-SUMIFS(资_净额,资_待补票,1,资_日期,"<="&{D})-SUM({hrng("已补票")}),2)', kind='⚠', fmtF=MONEY, link=SH_ADV,
       note='补到票在【发票登记】登一行，「垫付编号」填资金台帐那一行的编号；【个人垫付现金明细】右边「补票跟进」按天数列出来')
    it(S, '垫付编号', '老板垫付：编号没填或重复的行',
       H=('COUNTIFS(资_账户类型,5,资_有效,1)-COUNTIFS(资_账户类型,5,资_有效,1,资_编号键,"#*")'
          '+COUNTIF(资_提醒,"⚠ 垫付编号重复*")'), kind='⚠', link=SH_CASH,
       note='垫付账户每一笔都填编号（资金台帐第 3 行有「下一个垫付编号」），补票、还款都按编号对')
    it(S, '固定支出', '本月固定支出过了日子还没付清（条；F＝还差的金额）', H=f'SUMPRODUCT({fix})',
       F=f'ROUND(SUMPRODUCT({fix}*(固_每月金额-固_本月已付)),2)', kind='⚠', fmtF=MONEY, link=SH_FIX,
       note='付了以后资金台帐「编号」填固定支出的编号（或者收支类别、往来单位对得上）就算已付')
    it(S, '补票不抵', '补到专票但不抵进项（买设备、建厂费用、营业外支出，或类别不得抵扣、小规模）',
       H='COUNTIFS(票_补票行,">0",票_有效,1,票_补票额,0,票_税额,">0")', kind='ℹ', link=SH_INV,
       note='这些票只消掉「待补票」，进项税不抵（1601/1801/5711 补的票、不得抵扣的类别）；要抵的话付款时就选专票')
    it(S, '账户负', '现金、银行、微信、承兑账户：余额是负数的账户（截止日）', H=f'SUM({hrng("账户负")})',
       F=first_name('账户负', '账户_名称'), fmtF='@', kind='⚠', link=SH_DAY,
       note='钱不够还付出去了？多半漏记了收款、内部转进，或者【基础资料】② 账户期初余额没填')
    it(S, '截止早', '首页截止日期比最后一笔早（天）', H=f'IF(AND(ISNUMBER({HOME_C}),{D}<P_最后日期),P_最后日期-{D},0)',
       kind='⚠', link=SH_HOME, note='首页黄格填了截止日期，后面录的就不算进报表；看完了记得清空')
    it(S, '日期错', '日期打错的嫌疑（最后几笔比别的晚一个多月，截止日期没跟着走）',
       H=f'IF(AND(NOT(ISNUMBER({HOME_C})),P_截止被挡=1),P_近月笔数,0)', kind='⚠', link=SH_CASH,
       note='资金台帐、进货、销售、生产日报按日期从大到小排一下就找到，多半年份或月份填错了')
    return L


# ───────────────────────── 生成 ─────────────────────────
def positions(L):
    """版面：① 录入表 (第一行, 最后一行)、① 下面说明行、②③④ 每项的行 {键: 行}、每段标题行 [(段, 行)]、最后一行"""
    in0 = 8
    in1 = in0 + len(INPUTS) - 1
    note = in1 + 1
    r = note + 2
    pos, secs, cur = {}, [], None
    for d in L:
        if d['sec'] != cur:
            if cur is not None:
                r += 1
            cur = d['sec']
            secs.append((cur, r))
            r += 2
        pos[d['key']] = r
        r += 1
    return (in0, in1), note, pos, secs, r - 1


ROWS = {}                                  # 键 → 行号（model_check 用）


def _dn(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _link(ws, a, sheet, text=None):
    c = put(ws, a, text or f'{sheet} →', F_LINK, align=AC)
    link(c, sheet)


def build(wb, ctx=None):
    ws = wb[SH_CHK]
    widths(ws, {'A': 5, 'B': 22, 'C': 11, 'D': 11, 'E': 11, 'F': 13, 'G': 13, 'H': 13, 'I': 6, 'J': 13, 'K': 60})
    title(ws, '数 据 校 验', LAST, C_CHK,
          '★ 录完数看一眼这张：✗＝一定要改（不改报表会少算、算错），⚠＝看一下对不对、该跟进了，ℹ＝说明一下（不算问题），√＝没问题。'
          '先看 ① 每张录入表有几行问题（点表名过去，在最右边「这一行的问题」筛 ✗；「第一处」点了直接跳到那一行），再看 ② 基础资料、③ 账对不对、'
          '④ 该跟进的事。数都算到截止日期（首页黄格，空着＝最后一笔的日期）那个月。这张表全是自动的。')
    home_link(ws, 'J3')
    put(ws, 'B4', '要改的（✗）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('D4:E4')
    put(ws, 'D4', '要看的（⚠）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'G4', '="截止 "&TEXT(P_截止,"yyyy-mm-dd")&"（"&P_截止月&" 月）"', F_NOTE, align=AL, border=False)
    ws.row_dimensions[4].height = 26

    # 隐藏辅助列
    hf = helper_formulas()
    for k, n in HELPERS:
        c = HCOL[k]
        ws[f'{c}1'] = k
        ws[f'{c}1'].font = F_HELP
        for i in range(1, n + 1):
            ws[f'{c}{i + 1}'] = '=' + hf[k](i)
            ws[f'{c}{i + 1}'].font = F_HELP
        ws.column_dimensions[c].hidden = True
    ws.column_dimensions['L'].width = 3

    # ① 各录入表
    r = 6
    section(ws, r, SEC_C1, SEC_C2, '① 各录入表逐行检查（✗、⚠ 不含示例行；示例行另数）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '表（点名字过去）'), ('C', '✗ 要改\n（行）'), ('D', '⚠ 要看\n（行）'), ('E', '示例行'),
                       ('F', '右边没公式\n（行）'), ('G', '容量外有内容\n（格）'), ('H', '插了列\n（列）'), ('I', '结果'),
                       ('J', '第一处 ✗'), ('K', '常见原因、怎么改')], C_CHK, height=34)
    r1 = r + 2
    for i, (sh, p, note_c, chk_c, last, cap, why) in enumerate(INPUTS):
        rr = r1 + i
        end = chk_end(cap)
        nrng = f'{sh}!${note_c}${HDR}:${note_c}${end}'
        crng = f'{sh}!${chk_c}${HDR}:${chk_c}${end}'
        li = last_input(sh)
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        link(put(ws, f'B{rr}', sh, F_LINK, align=AL), sh)
        put(ws, f'C{rr}', f'=MAX(0,COUNTIF({p}_错,"✗*")-COUNTIFS({nrng},"示例*",{crng},"✗*"))', F_AUTOB, fmt=INT, align=AC)
        put(ws, f'D{rr}', (f'=MAX(0,COUNTIF({p}_提醒,"⚠*")-COUNTIFS({p}_错,"✗*",{p}_提醒,"⚠*")'
                           f'-COUNTIFS({nrng},"示例*",{crng},"⚠*"))'), F_AUTOB, fmt=INT, align=AC)
        put(ws, f'E{rr}', f'=COUNTIF({nrng},"示例*")', F_AUTO, fmt=INT, align=AC)
        put(ws, f'F{rr}', f'=MAX(0,COUNTIF({p}_末行,">0")-COUNTIF({crng},"?*")+COUNTIF({sh}!${chk_c}${HDR},"?*"))',
            F_AUTOB, fmt=INT, align=AC)
        put(ws, f'G{rr}', '=' + over_rows(sh, HDR, cap, li), F_AUTOB, fmt=INT, align=AC)
        put(ws, f'H{rr}', f'=COLUMNS({sh}!$A${HDR}:${last}${HDR})-{CI(last)}', F_AUTOB, fmt='+0;-0;"-"', align=AC)
        put(ws, f'I{rr}', (f'=IF(OR(C{rr}>0,G{rr}>0,H{rr}<>0),"✗",IF(OR(D{rr}>0,F{rr}>0),"⚠",IF(E{rr}>0,"ℹ","√")))'),
            F_TXTB, align=AC)
        fr = f'INDEX({p}_录入行,MATCH("✗*",{p}_错,0))'
        put(ws, f'J{rr}', f'=IFERROR(IF(C{rr}>0,HYPERLINK("#\'{sh}\'!A"&{fr},"第 "&{fr}&" 行 →"),""),"")', F_LINK, align=AC)
        put(ws, f'K{rr}', why + '。示例行（灰底、备注写「示例」）启用前整行删掉', F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 42
        ROWS[sh] = rr
    rr = r1 + len(INPUTS)
    ws.merge_cells(f'B{rr}:K{rr}')
    put(ws, f'B{rr}', ('「右边没公式」：插进来的行最右边几格是空的（不影响记账，只是那一行看不到问题）：选上一行右边灰色几格往下拉一下就有了。'
                       '「容量外有内容」：超过容量的行不算（找我加容量或另起一本）。「插了列」：录入表的列不能插、不能删，撤销或删掉/补回来。'),
        F_NOTE, align=ALW, border=False)
    ws.row_dimensions[rr].height = 30
    in_rows = (r1, r1 + len(INPUTS) - 1)

    # ②③④
    L = items(wb)
    _ins, _note, pos, secs, last_row = positions(L)
    assert _ins == in_rows and _note == rr, (_ins, in_rows, _note, rr)
    titles = {2: '② 基础资料（建账日、期初、名称重复、容量）', 3: '③ 账对不对（截止日期那个月月末；报表的核对格按报表上选的月份）',
              4: '④ 该跟进的事（截至截止日期）'}
    hdr = [('A', '#'), ('B', '检查'), ('F', '数 1'), ('G', '数 2'), ('H', '数量／差'), ('I', '结果'), ('J', '去哪改'),
           ('K', '说明（怎么改）')]
    for sec, sr in secs:
        section(ws, sr, SEC_C1, SEC_C2, titles[sec], C_CHK)
        header(ws, sr + 1, hdr, C_CHK, height=24)
        ws.merge_cells(f'B{sr + 1}:E{sr + 1}')
    first_item = secs[0][1] + 2
    no, cur = 0, None
    for d in L:
        r = pos[d['key']]
        if d['sec'] != cur:
            cur, no = d['sec'], 0
        no += 1
        put(ws, f'A{r}', no, F_TXT, align=AC)
        ws.merge_cells(f'B{r}:E{r}')
        lab = d['label']
        put(ws, f'B{r}', lab, F_TXT, align=ALW)
        for col in 'CDE':
            ws[f'{col}{r}'].border = BD
        T = None
        if d['text']:
            ws.merge_cells(f'F{r}:H{r}')
            T = f'F{r}'
            put(ws, T, '=' + d['text'], F_TXT, align=ALW)
            for col in 'GH':
                ws[f'{col}{r}'].border = BD
        else:
            for col, key, fm in (('F', 'F', d['fmtF']), ('G', 'G', d['fmtG']), ('H', 'H', d['fmt'])):
                v = d[key]
                if v is None:
                    put(ws, f'{col}{r}', None, F_TXT)
                    continue
                v = v.replace('{F}', f'F{r}').replace('{G}', f'G{r}')
                put(ws, f'{col}{r}', '=' + v, F_AUTOB if col == 'H' else F_AUTO, fmt=fm,
                    align=AL if fm == '@' else AR)
        cond = d['cond'].replace('{H}', f'H{r}').replace('{F}', f'F{r}').replace('{T}', T or '""')
        kind = d['kind']
        kind_f = kind.replace('{T}', T or '""') if kind.startswith('IF(') else f'"{kind}"'
        put(ws, f'I{r}', f'=IF({cond},{kind_f},"√")', F_TXTB, align=AC)
        if d['link']:
            _link(ws, f'J{r}', d['link'])
        put(ws, f'K{r}', d['note'], F_NOTE, align=ALW)
        ws.row_dimensions[r].height = 42 if len(d['note']) > 60 else 30
        ROWS[d['key']] = r
    xs = (f'=SUM(C{in_rows[0]}:C{in_rows[1]})+COUNTIF(G{in_rows[0]}:G{in_rows[1]},">0")'
          f'+COUNTIF(H{in_rows[0]}:H{in_rows[1]},"<>0")+COUNTIF(I{first_item}:I{last_row},"✗")')
    wsum = (f'=SUM(D{in_rows[0]}:D{in_rows[1]})+COUNTIF(F{in_rows[0]}:F{in_rows[1]},">0")'
            f'+COUNTIF(I{first_item}:I{last_row},"⚠")')
    put(ws, CHK_X, xs, F_KPI_V, fill('FFFFFFFF'), '0', AC)
    put(ws, CHK_W, wsum, F_KPI_V, fill('FFFFFFFF'), '0', AC)
    q = SH_CHK
    _dn(wb, '校验_错数', f'{q}!${CHK_X[0]}${CHK_X[1:]}')
    _dn(wb, '校验_警数', f'{q}!${CHK_W[0]}${CHK_W[1:]}')
    for nm, key, col in (('负库存', '负库存', 'H'), ('工资没登', '工资没登', 'H'), ('待定价进货笔数', '待定价进', 'H'),
                         ('待定价进货金额', '待定价进', 'F'), ('待补票笔数', '待补票', 'H'), ('待补票金额', '待补票', 'F'),
                         ('固定支出没付条数', '固定支出', 'H'), ('固定支出没付金额', '固定支出', 'F')):
        _dn(wb, f'校验_{nm}', f'{q}!${col}${ROWS[key]}')

    # 条件格式
    red = Font(name=YH, sz=10, bold=True, color='FFC00000')
    rng = f'I{in_rows[0]}:I{last_row}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'I{in_rows[0]}="✗"'], font=red, fill=cf_bg('FFFFC7CE')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'I{in_rows[0]}="⚠"'], font=F_WARN, fill=cf_bg('FFFFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'I{in_rows[0]}="ℹ"'], font=F_INFO))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'I{in_rows[0]}="√"'], font=F_OK))
    for col in 'CG':
        a = f'{col}{in_rows[0]}:{col}{in_rows[1]}'
        ws.conditional_formatting.add(a, FormulaRule(formula=[f'{col}{in_rows[0]}>0'], font=red))
    for col in 'DF':
        a = f'{col}{in_rows[0]}:{col}{in_rows[1]}'
        ws.conditional_formatting.add(a, FormulaRule(formula=[f'{col}{in_rows[0]}>0'], font=F_WARN))
    a = f'H{in_rows[0]}:H{in_rows[1]}'
    ws.conditional_formatting.add(a, FormulaRule(formula=[f'H{in_rows[0]}<>0'], font=red))
    ws.conditional_formatting.add(CHK_X, FormulaRule(formula=[f'{CHK_X}=0'], font=Font(name=YH, sz=13, bold=True, color='FF548235')))
    ws.conditional_formatting.add(CHK_W, FormulaRule(formula=[f'{CHK_W}=0'], font=Font(name=YH, sz=13, bold=True, color='FF548235')))
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:4', landscape=True)
    ws.print_area = f'A1:{LAST}{last_row}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
