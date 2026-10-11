# -*- coding: utf-8 -*-
"""A072 化工 · 记账：隐藏表 _分（记账分录槽位）、_凭（凭证排序键）、_矩（科目×月矩阵）；
   查看表 凭证汇总、记账凭证（打印，照 A062 m1008：A4 纵向一张两页、75%、中缝）、科目余额表、年末结转。

   _分（layout.FEN_BLOCKS）：A 来源 B n（月末块＝月）C 腿 写常量；D 科目号 E 借 F 贷 G 累 H 键 每块每列一个公式（按腿号参数化）。
     每腿先算带符号的借方额 x：借＝MAX(x,0)，贷＝MAX(−x,0)；月末块只在 P_建账月≤月≤P_截止月 出金额。
     G 累＝IF(腿=1,0,上一行 G)＋这一腿有没有额；H 键＝有额：科目号×100＋月，否则 0。
     费用归属（21）的 224103 腿、补票进项（22）的 222101 腿＝本张其他腿的轧差（借贷一定平；目标科目不在费用明细清单里的那部分记不进来，224103 留余额）。
   _凭：A 排序键＝IF(这张凭证有分录, INT(日期)×10^7＋来源×10^4＋n（月末＝月）, "")；B 来源、C n 常量。
   凭证定位一律按位置：起点＝FEN_BASE−1＋(n−1)×腿数（相对 分_ 名称第 1 行），条数＝INDEX(分_累,终点)，
     合计＝SUM(INDEX(分_借,起):INDEX(分_借,终))，第 i 条＝起点−1＋MATCH(i, INDEX(分_累,起):INDEX(分_累,终), 0)；
     手工凭证在手工块 300 行里 MATCH 手_印键。

   私有（别的模块不要引用；要用就用本模块加的名称）：
   _分 J2＝员工保险类别的付号（没有＝224199）。
   _凭 E..H 第 2..17 行：来源代码、起点基（FEN_BASE−1）、腿数、名称。
   _矩 AA 排序键（编码补到 6 位×1000＋科目号，给科目余额表、年末结转排顺序）；AB..AM 汇借、AN..AY 汇贷、AZ 汇期初
       （一级科目＝下级合计＋自己，二级＝自己）；BA 年末（自己，全年）；BB 全年借；BC 全年贷；BD 汇年末。
       名称：矩_排序键、矩_汇借、矩_汇贷、矩_汇期初、矩_年末、矩_汇年末。
   凭证汇总 M..R 隐藏（键、来源号、n、起点、终点、起页键），T/U 隐藏标量；名称 汇总_*（记账凭证用）。
   记账凭证 L..R 隐藏：每页第 1 行 L k、M 本凭证第几页、N 来源、O n、P 起点、Q 终点、R 页数；分录行 L＝分_ 第几行。
   年末结转 X 列以右隐藏辅助表（出借人、⑦ 槽位、⑧ 槽位）。
   ⑧：1801、1123 按摊销来源一行一个（资金台帐 分摊＝1/2 的每一笔、⑧ 期初 1801/1123 每一行），按截止月算还剩几个月、还剩多少、摊进哪
     （1801 写 制造/管理，1123 写「编码 名称」），新本 ⑧ 接着摊；1123 摊不动的 ⑧ 行原样带过去（⚠）；两个家族各剩一行「其余」
     ＝矩_汇年末（含下级）－各来源剩余（手工凭证等记的）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, Alignment, Border, Side, Protection
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break, RowBreak
from openpyxl.worksheet.protection import SheetProtection
from common import *
from layout import *                      # layout 的 norm/esc 盖过 common 的
from s_mirror import unit_match

NP, NF, NOH = N_PRD, N_FEE, N_OH
KM = KEY_M                                # 10^7
VM = VID_M                                # 10^4
B4 = FEN_BASE[SRC_MJ] - 2                 # 手工第 n 条在 分_ 名称第 B4＋n 个
ME_NAME = {c: nm for c, nm, _l in ME_TYPES}
SRC_NAME = {SRC_CASH: '资金台帐', SRC_PUR: '进货', SRC_SALE: '销售', SRC_MJ: '手工凭证', **ME_NAME}
SRC_LIST = [s for s, _r, _n, _l in FEN_BLOCKS]
TB_SRC = f"{H_PZ}!$E$2:$E${len(SRC_LIST) + 1}"
TB_BASE = f"{H_PZ}!$F$2:$F${len(SRC_LIST) + 1}"
TB_LEGS = f"{H_PZ}!$G$2:$G${len(SRC_LIST) + 1}"
TB_NAME = f"{H_PZ}!$H$2:$H${len(SRC_LIST) + 1}"
INS = f'{H_FEN}!$J$2'
LINES = 8                                 # 每页凭证 8 条分录
VP_ROWS = 14                              # 每页凭证 14 行
YELLOW = FILL_SEL


def K(code):
    return kh(code)


def ix(nm, i):
    return f'INDEX({nm},{i})'


def _dn(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def _lab(ws, a, t):
    ws[a] = t
    ws[a].font = F_HELP


def protect(ws):
    """工作表保护：黄格不锁，无密码；允许筛选、调列宽行高，不准排序"""
    for row in ws.iter_rows():
        for c in row:
            if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb in ('FFFFFF00', 'FFFFF7E0'):
                c.protection = Protection(locked=False)
    ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=True, formatColumns=False, formatRows=False,
                                    formatCells=False, selectLockedCells=False, selectUnlockedCells=False)


def month_raw(c):
    """月份格 → 1～12；空＝-1；看不懂或不是本年度＝0（认 9、9月、2026-9、2026年9月、202609、日期）"""
    v = f'TRIM({c}&"")'
    t = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({v},"年","-"),"月",""),"/","-"),".","-")'
    tv = (f'IFERROR(IF(ISNUMBER(FIND("-",{t})),IF(VALUE(LEFT({t},4))=P_年度,VALUE(MID({t},FIND("-",{t})+1,2)),0),'
          f'VALUE({t})),0)')
    nv = (f'IF(AND({c}>=1,{c}<=12),{c},IF(AND({c}>=200001,{c}<=209912),IF(INT({c}/100)=P_年度,MOD({c},100),0),'
          f'IF({c}>=36526,IF(YEAR({c})=P_年度,MONTH({c}),0),0)))')
    return f'IF({v}="",-1,IF(ISNUMBER({c}),{nv},{tv}))'


# ═══════════════════════════════ _分 ═══════════════════════════════
def fen_expr(src, legs, r):
    """第 r 行（来源 src）的 (科目号, 带符号借方额 x, 月) 公式片段"""
    Bn, C = f'$B{r}', f'$C{r}'
    if src == SRC_CASH:
        ne = ix('资_净额', Bn)
        d = f'CHOOSE({C},{ix("资_账户科目号", Bn)},{ix("资_对方号", Bn)},{K("222101")})'
        x = f'CHOOSE({C},{ne},-SIGN({ne})*{ix("资_不含税", Bn)},-SIGN({ne})*{ix("资_税额", Bn)})'
        return d, x, ix('资_月', Bn)
    if src == SRC_PUR:
        d = f'CHOOSE({C},{ix("进_借方号", Bn)},{K("222101")},{K("2202")},{K("2202")})'
        x = (f'CHOOSE({C},{ix("进_入账额", Bn)},{ix("进_货税", Bn)}+{ix("进_运税", Bn)},-{ix("进_货款计", Bn)},'
             f'-{ix("进_运费计", Bn)})')
        return d, x, ix('进_月', Bn)
    if src == SRC_SALE:
        d = (f'CHOOSE({C},{K("1122")},IF({ix("销_出库", Bn)}=2,{K("5051")},{K("5001")}),{K("222102")},'
             f'IF(P_运费进成本=1,{K("540102")},{K("560101")}),{K("222101")},{K("2202")})')
        x = (f'CHOOSE({C},{ix("销_收入含税", Bn)},-{ix("销_收入", Bn)},-{ix("销_销项", Bn)},{ix("销_运费入账", Bn)},'
             f'{ix("销_运税", Bn)},-{ix("销_运费计", Bn)})')
        return d, x, ix('销_月', Bn)
    if src == SRC_MJ:
        d = ix('手_科目号', Bn)
        x = f'({ix("手_借", Bn)}-{ix("手_贷", Bn)})*{ix("手_记账", Bn)}'
        return d, x, ix('手_月', Bn)
    m, g = Bn, C                                            # 月末：B＝月，C＝腿
    col = lambda nm: f'INDEX({nm},0,{m})'                 # 某字段块第 m 月那一列
    at = lambda nm, i: f'INDEX({nm},{i},{m})'
    if src == 11:      # 领料
        d = f'IF({g}<={NP},{K("400101")},{K("1403")})'
        x = f'IF({g}<={NP},{at("本_DM", g)},-SUM({col("本_DM")}))'
    elif src == 12:    # 废料回收
        d = f'IF({g}=1,{K("1403")},{K("400101")})'
        x = f'IF({g}=1,SUM({col("本_废料扣")}),-{at("本_废料扣", f"{g}-1")})'
    elif src == 13:    # 制造费用分配
        d = f'IF({g}<={NP},{K("400103")},INDEX(制费明细_号,{g}-{NP}))'
        x = f'IF({g}<={NP},{at("本_OH", g)},-{at("本池_本月分配", f"{g}-{NP}")})'
    elif src == 14:    # 完工入库
        k3 = f'INT(({g}-1)/{NP})+1'
        p = f'MOD({g}-1,{NP})+1'
        d = f'CHOOSE({k3},{K("1405")},{K("400101")},{K("400103")})'
        x = f'CHOOSE({k3},{at("本_完工", p)},-{at("本_转出DM", p)},-{at("本_OH", p)})'
    elif src == 15:    # 销售成本
        d = f'CHOOSE(MIN({g},4),{K("540101")},{K("5402")},{K("1403")},{K("1405")})'
        x = (f'CHOOSE(MIN({g},4),SUM({col("品_销额")}),SUM({col("料_销额")}),-SUM({col("料_销额")}),'
             f'-{at("品_销额", f"MAX(1,{g}-3)")})')
    elif src == 16:    # 盘点调整
        yr, yp = f'SUM({col("料_盘额")})', f'SUM({col("品_盘额")})'
        d = f'CHOOSE({g},{K("1403")},{K("410199")},{K("1405")},{K("5711")},{K("5301")})'
        x = f'CHOOSE({g},{yr},-{yr},{yp},MAX(-{yp},0),-MAX({yp},0))'
    elif src == 17:    # 折旧
        blk = col('卡_月折旧块')
        d = f'CHOOSE({g},{K("410107")},{K("560206")},{K("560199")},{K("1602")})'
        x = f'IF({g}<=3,SUMIF(卡_部门号,{g},{blk}),-SUM({blk}))'
    elif src == 18:    # 摊销
        d = f'IF({g}<={NF},INDEX(费用明细_号,{g}),IF({g}={NF + 1},{K("1801")},{K("1123")}))'
        x = (f'IF({g}<={NF},{at("摊_摊1801", g)}+{at("摊_摊1123", g)},'
             f'IF({g}={NF + 1},-INDEX(摊_贷1801,1,{m}),-INDEX(摊_贷1123,1,{m})))')
    elif src == 19:    # 工资计提
        s = lambda f: f'SUM({col("工汇_" + f)})'
        d = (f'CHOOSE({g},{K("410104")},{K("410105")},{K("560201")},{K("560103")},{K("2211")},{K("2211")},{INS},'
             f'{K("2211")},{K("5301")})')
        x = (f'CHOOSE({g},{at("工汇_应发", 1)},{at("工汇_应发", 2)},{at("工汇_应发", 3)},{at("工汇_应发", 4)},'
             f'-{s("应发")},{s("扣社保")},-{s("扣社保")},{s("其他扣款")},-{s("其他扣款")})')
    elif src == 20:    # 利息计提
        a, b = f'SUM({col("息_借本月利息")})', f'SUM({col("息_投本月利息")})'
        d = f'CHOOSE({g},{K("560301")},{K("2231")},{K("310401")},{K("2232")})'
        x = f'CHOOSE({g},{a},-{a},{b},-{b})'
    elif src == 21:    # 费用归属：224103 腿＝本张前面各腿轧差
        no = f'INDEX(费用明细_号,{g})'
        d = f'IF({g}<={NF},{no},{K("224103")})'
        x = (f'IF({g}<={NF},IF({no}=0,0,SUMIFS(资_费用额,资_跨月,1,资_原对方号,{no},资_费用月,{m})),'
             f'SUM(F{r - NF}:F{r - 1})-SUM(E{r - NF}:E{r - 1}))')
    elif src == 22:    # 补票进项：222101 腿＝本张后面各腿轧差
        no = f'INDEX(费用明细_号,{g}-1)'
        d = f'IF({g}=1,{K("222101")},{no})'
        x = (f'IF({g}=1,SUM(F{r + 1}:F{r + NF})-SUM(E{r + 1}:E{r + NF}),'
             f'IF({no}=0,0,-SUMIFS(票_补票额,票_补票目标号,{no},票_月,{m})))')
    else:
        raise ValueError(src)
    x = f'IF(AND({m}>=P_建账月,{m}<=P_截止月),{x},0)'
    return d, x, m


def build_fen(ws):
    for k, c in FEN_COLS.items():
        _lab(ws, f'{c}1', k)
    _lab(ws, 'J1', '员工保险付号')
    ins = 'IFERROR(INDEX(类别_付号,MATCH("员工保险",类别_名称,0)),0)'
    ws['J2'] = f'=IF({ins}>0,{ins},{K("224199")})'
    for src, row0, cnt, legs in FEN_BLOCKS:
        for n in range(1, cnt + 1):
            for g in range(1, legs + 1):
                r = row0 + (n - 1) * legs + g - 1
                ws[f'A{r}'] = src
                ws[f'B{r}'] = n
                ws[f'C{r}'] = g
                d, x, mo = fen_expr(src, legs, r)
                ws[f'D{r}'] = f'={d}'
                ws[f'E{r}'] = f'=ROUND(MAX({x},0),2)'
                ws[f'F{r}'] = f'=ROUND(MAX(-({x}),0),2)'
                ws[f'G{r}'] = f'=IF($C{r}=1,0,G{r - 1})+IF(OR(E{r}<>0,F{r}<>0),1,0)'
                ws[f'H{r}'] = f'=IF(OR(E{r}<>0,F{r}<>0),D{r}*100+{mo},0)'
                ws[f'E{r}'].number_format = ws[f'F{r}'].number_format = '#,##0.00'


# ═══════════════════════════════ _凭 ═══════════════════════════════
def build_pz(ws):
    _lab(ws, 'A1', '排序键')
    _lab(ws, 'B1', '来源')
    _lab(ws, 'C1', 'n')
    for c, t in zip('EFGH', ('来源', '起点基', '腿数', '名称')):
        _lab(ws, f'{c}1', t)
    for i, src in enumerate(SRC_LIST):
        r = i + 2
        ws[f'E{r}'] = src
        ws[f'F{r}'] = FEN_BASE[src] - 1
        ws[f'G{r}'] = FEN_LEGS[src]
        ws[f'H{r}'] = SRC_NAME[src]
    dt = {SRC_CASH: '资_日期', SRC_PUR: '进_日期', SRC_SALE: '销_日期'}
    for src, row0, cnt in PZ_BLOCKS:
        legs, fb = FEN_LEGS[src], FEN_BASE[src]
        for n in range(1, cnt + 1):
            r = row0 + n - 1
            ws[f'B{r}'] = src
            ws[f'C{r}'] = n
            C = f'$C{r}'
            end = f'INDEX(分_累,{fb - 2}+{C}*{legs})'
            if src in dt:
                f = f'=IF({end}>0,INT(INDEX({dt[src]},{C}))*{KM}+{src * VM}+{C},"")'
            elif src == SRC_MJ:
                f = (f'=IF(AND(INDEX(手_记账,{C})=1,INDEX(手_组首,{C})={C}),'
                     f'INT(INDEX(手_日期,{C}))*{KM}+{src * VM}+{C},"")')
            else:
                f = (f'=IF(AND({C}>=P_建账月,{C}<=P_截止月),IF({end}>0,DATE(P_年度,{C}+1,0)*{KM}+{src * VM}+{C},""),"")')
            ws[f'A{r}'] = f
            ws[f'A{r}'].number_format = '0'


# ═══════════════════════════════ _矩 ═══════════════════════════════
JU_PRIV = dict(排序键='AA', 汇借='AB', 汇贷='AN', 汇期初='AZ', 年末='BA', 全年借='BB', 全年贷='BC', 汇年末='BD')
JR = N_COA + 1


def jcol(c0, m):
    return CL(CI(c0) + m - 1)


def build_ju(wb, ws):
    _lab(ws, 'A1', '科目号')
    for m in MONTHS:
        _lab(ws, f'{jcol(JU_DR0, m)}1', f'借{m}')
        _lab(ws, f'{jcol(JU_CR0, m)}1', f'贷{m}')
        _lab(ws, f'{jcol(JU_PRIV["汇借"], m)}1', f'汇借{m}')
        _lab(ws, f'{jcol(JU_PRIV["汇贷"], m)}1', f'汇贷{m}')
    _lab(ws, f'{JU_OPEN}1', '期初')
    for k in ('排序键', '汇期初', '年末', '全年借', '全年贷', '汇年末'):
        _lab(ws, f'{JU_PRIV[k]}1', k)
    P = JU_PRIV
    for r in range(1, N_COA + 1):
        w = r + 1
        A = f'$A{w}'
        ws[f'A{w}'] = r
        lv1 = f'INDEX(科目_级次,{A})=1'
        for m in MONTHS:
            for c0, side, h0 in ((JU_DR0, '借', P['汇借']), (JU_CR0, '贷', P['汇贷'])):
                c = jcol(c0, m)
                ws[f'{c}{w}'] = f'=SUMIF(分_键,{A}*100+{m},分_{side})'
                ws[f'{jcol(h0, m)}{w}'] = f'=IF({lv1},SUMIF(科目_一级号,{A},{c}$2:{c}${JR}),{c}{w})'
        ws[f'{JU_OPEN}{w}'] = f'=INDEX(科目_期初,{A})'
        ws[f'{P["排序键"]}{w}'] = (f'=IF(INDEX(科目_有效,{A})=1,VALUE(LEFT(INDEX(科目_编码,{A})&"00",6))*1000+{A},"")')
        ws[f'{P["汇期初"]}{w}'] = f'=IF({lv1},SUMIF(科目_一级号,{A},${JU_OPEN}$2:${JU_OPEN}${JR}),{JU_OPEN}{w})'
        d1, d2 = f'{JU_DR0}{w}:{jcol(JU_DR0, 12)}{w}', f'{JU_CR0}{w}:{jcol(JU_CR0, 12)}{w}'
        ws[f'{P["全年借"]}{w}'] = f'=SUM({d1})'
        ws[f'{P["全年贷"]}{w}'] = f'=SUM({d2})'
        ws[f'{P["年末"]}{w}'] = f'=ROUND({JU_OPEN}{w}+{P["全年借"]}{w}-{P["全年贷"]}{w},2)'
        h1 = f'{P["汇借"]}{w}:{jcol(P["汇借"], 12)}{w}'
        h2 = f'{P["汇贷"]}{w}:{jcol(P["汇贷"], 12)}{w}'
        ws[f'{P["汇年末"]}{w}'] = f'=ROUND({P["汇期初"]}{w}+SUM({h1})-SUM({h2}),2)'
    rng = lambda c, c2=None: f'{H_JU}!${c}$2:${c2 or c}${JR}'
    _dn(wb, '矩_排序键', rng(P['排序键']))
    _dn(wb, '矩_汇借', rng(P['汇借'], jcol(P['汇借'], 12)))
    _dn(wb, '矩_汇贷', rng(P['汇贷'], jcol(P['汇贷'], 12)))
    _dn(wb, '矩_汇期初', rng(P['汇期初']))
    _dn(wb, '矩_年末', rng(P['年末']))
    _dn(wb, '矩_汇年末', rng(P['汇年末']))
    _dn(wb, '矩_全年借', rng(P['全年借']))
    _dn(wb, '矩_全年贷', rng(P['全年贷']))


# ═══════════════════════════════ 凭证汇总 ═══════════════════════════════
VL = dict(no='A', date='B', memo='C', dr='D', cr='E', ok='F', lines='G', pages='H', page0='I', a4='J', src='K',
          key='M', sno='N', n='O', st='P', en='Q', pk='R')
VL_LAST = 'K'
VL_HDR, VL_R0 = 4, 5
VL_R1 = VL_R0 + N_VLIST - 1
VL_SC = ['输入', '认出', '选月', '以前', '本月', '显示', '总页数', 'A4', '不平']


def vsc(k):
    return f'$U${3 + VL_SC.index(k)}'


def vrng(k):
    c = VL[k]
    return f"{SH_VLIST}!${c}${VL_R0}:${c}${VL_R1}"


def memo_expr(N, O):
    """凭证汇总摘要（N＝来源号格、O＝n 格）"""
    lk, pk = f'INDEX(进_料号,{O})', f'INDEX(进_品号,{O})'
    nm_j = f'IF({lk}>0,INDEX(料档_名称,{lk}),IF({pk}>0,INDEX(品档_名称,{pk}),""))'
    q = f'INDEX(进_数量,{O})'
    unit = f'IF({lk}>0,INDEX(料档_单位,{lk}),"吨")'
    jin = f'"进货 "&{nm_j}&" "&IF({q}=0,"扣款/折让",ROUND({q},3)&{unit})&" "&INDEX(进_供应商,{O})'
    sp, sl = f'INDEX(销_品号,{O})', f'INDEX(销_料号,{O})'
    nm_s = f'IF({sp}>0,INDEX(品档_名称,{sp}),IF({sl}>0,INDEX(料档_名称,{sl}),""))'
    t = f'INDEX(销_吨位,{O})'
    xs = f'"销售 "&INDEX(销_客户,{O})&" "&{nm_s}&" "&IF({t}=0,"扣款/折让",ROUND({t},3)&"吨")'
    mj = f'INDEX({inref(SH_MJ, MJ_COLS["摘要"])},{O}+1)&""'
    me = f'{O}&"月"&INDEX({TB_NAME},MATCH({N},{TB_SRC},0))'
    return (f'IF({N}={SRC_CASH},INDEX(资_显示摘要,{O})&"",IF({N}={SRC_PUR},{jin},IF({N}={SRC_SALE},{xs},'
            f'IF({N}={SRC_MJ},{mj},{me}))))')


def build_vlist(wb, ws):
    c = VL
    widths(ws, {c['no']: 9, c['date']: 11, c['memo']: 46, c['dr']: 16, c['cr']: 16, c['ok']: 13, c['lines']: 7,
                c['pages']: 6, c['page0']: 8, c['a4']: 10, c['src']: 14})
    title(ws, '凭 证 汇 总', VL_LAST, C_ACC,
          '★ 选一个月（B3 黄格，空着＝截止日期那个月），列出这个月全部凭证：资金台帐一行一张、进货一车一张、销售一车一张、手工凭证一组一张，'
          '月末自动的领料、废料回收、制造费用分配、完工入库、销售成本、盘点调整、折旧、摊销、工资计提、利息计提、费用归属、补票进项各一张（月末最后一天）。'
          '凭证号＝当月第几号（按日期排，同一天按 资金→进货→销售→手工→月末 的顺序）。一张凭证超过 8 条分录接下一页；「打印起始页」是它在【记账凭证】里从第几页开始，'
          '一张 A4 纸上下印两页凭证：打印页码范围填 1 到右上角「A4 张数」；只印某一张，看它「在 A4 第几张」。凭证号印不印看【基础资料】参数「凭证号要不要印出来」。'
          '只能筛选、不能排序（会计组的表都锁了，黄格能改）。')
    ws.row_dimensions[3].height = 26
    for i in range(1, CI(VL_LAST) + 1):
        put(ws, f'{CL(i)}3', None, fill_=fill('FFDDEBF7'), border=False)
    selector(ws, 'A3', '月份', 'B3', None, '"1,2,3,4,5,6,7,8,9,10,11,12"',
             prompt='选 1～12（本年度）；空着＝截止日期那个月；填 2026-9、9月 也认')
    S = vsc
    f = {
        '输入': '=' + month_raw('$B$3'),
        '认出': f'=IF({S("输入")}=-1,1,IF(AND({S("输入")}>=1,{S("输入")}<=12,INT({S("输入")})={S("输入")}),1,0))',
        '选月': f'=IF(OR({S("输入")}=-1,{S("认出")}=0),P_截止月,{S("输入")})',
        '以前': f'=COUNTIF(凭_键,"<"&DATE(P_年度,{S("选月")},1)*{KM})',
        '本月': f'=COUNTIF(凭_键,"<"&DATE(P_年度,{S("选月")}+1,1)*{KM})-{S("以前")}',
        '显示': f'=MIN({S("本月")},{N_VLIST})',
        '总页数': f'=SUM(${c["pages"]}${VL_R0}:${c["pages"]}${VL_R1})',
        'A4': f'=ROUNDUP({S("总页数")}/2,0)',
        '不平': f'=COUNTIF(${c["ok"]}${VL_R0}:${c["ok"]}${VL_R1},"✗*")',
    }
    for i, k in enumerate(VL_SC):
        _lab(ws, f'T{3 + i}', k)
        ws[f'U{3 + i}'] = f[k]
        ws[f'U{3 + i}'].font = F_HELP
    ws.merge_cells('C3:D3')
    put(ws, 'C3', (f'="实际："&P_年度&" 年 "&{S("选月")}&" 月"&IF(TRIM($B$3&"")="","（空着＝截止日期那个月）","")'
                   f'&IF({S("认出")}=0,"　⚠ 月份没认出来或不是本年度，先按截止日期那个月","")'), F_NOTE, align=ALW, border=False)
    kp = [('E3', '凭证张数：', 'F3', f'={S("本月")}', '0" 张"'), ('G3', '凭证页数：', 'H3', f'={S("总页数")}', '0" 页"'),
          ('I3', 'A4 张数：', 'J3', f'={S("A4")}', '0" 张"')]
    for lc, lab, vcl, fm, fmt in kp:
        put(ws, lc, lab, F_KPI_L, fill('FFDDEBF7'), align=AR, border=False)
        put(ws, vcl, fm, Font(name=YH, sz=12, bold=True, color='FFC00000'), fill('FFDDEBF7'), fmt, AL, border=False)
    put(ws, 'K3', (f'=IF({S("本月")}>{N_VLIST},"✗ 本月 "&{S("本月")}&" 张，只列前 {N_VLIST} 张",'
                   f'IF({S("不平")}>0,"✗ "&{S("不平")}&" 张借贷不平",IF({S("总页数")}>{N_VPAGE},"超过 {N_VPAGE} 页：记账凭证分批印",'
                   f'"")))'), F_RED, fill('FFDDEBF7'), align=AL, border=False)
    header(ws, VL_HDR, [(c['no'], '凭证号'), (c['date'], '日期'), (c['memo'], '摘要'), (c['dr'], '借方合计'),
                        (c['cr'], '贷方合计'), (c['ok'], '借贷平衡'), (c['lines'], '分录\n条数'), (c['pages'], '页数'),
                        (c['page0'], '打印\n起始页'), (c['a4'], '在 A4\n第几张'), (c['src'], '来源')], C_ACC, height=34)
    for k, t in (('key', '键'), ('sno', '来源号'), ('n', 'n'), ('st', '起点'), ('en', '终点'), ('pk', '起页键')):
        _lab(ws, f'{c[k]}{VL_HDR}', t)
    blk_e = f'{H_FEN}!$E${FEN_BASE[SRC_MJ]}:$E${FEN_BASE[SRC_MJ] + N_MJ - 1}'
    blk_f = f'{H_FEN}!$F${FEN_BASE[SRC_MJ]}:$F${FEN_BASE[SRC_MJ] + N_MJ - 1}'
    for k in range(1, N_VLIST + 1):
        r = VL_R0 + k - 1
        X = lambda nm: f'${c[nm]}{r}'
        M, N, O, P, Q = X('key'), X('sno'), X('n'), X('st'), X('en')
        nil = f'{M}=""'
        kk = f'(ROW()-{VL_HDR})'
        ws[M[1:]] = f'=IF({kk}>{S("显示")},"",SMALL(凭_键,{S("以前")}+{kk}))'
        ws[f'{c["no"]}{r}'] = f'=IF({nil},"",{kk})'
        ws[f'{c["date"]}{r}'] = f'=IF({nil},"",INT({M}/{KM}))'
        ws[N[1:]] = f'=IF({nil},"",INT(MOD({M},{KM})/{VM}))'
        ws[O[1:]] = f'=IF({nil},"",MOD({M},{VM}))'
        t = f'MATCH({N},{TB_SRC},0)'
        ws[P[1:]] = f'=IF({nil},"",INDEX({TB_BASE},{t})+({O}-1)*INDEX({TB_LEGS},{t}))'
        ws[Q[1:]] = f'=IF({nil},"",{P}+INDEX({TB_LEGS},{t})-1)'
        ws[f'{c["lines"]}{r}'] = (f'=IF({nil},"",IF({N}={SRC_MJ},COUNTIFS(手_组首,{O},手_记账,1),INDEX(分_累,{Q})))')
        ws[f'{c["dr"]}{r}'] = (f'=IF({nil},"",IF({N}={SRC_MJ},SUMIFS({blk_e},手_组首,{O},手_记账,1),'
                               f'SUM(INDEX(分_借,{P}):INDEX(分_借,{Q}))))')
        ws[f'{c["cr"]}{r}'] = (f'=IF({nil},"",IF({N}={SRC_MJ},SUMIFS({blk_f},手_组首,{O},手_记账,1),'
                               f'SUM(INDEX(分_贷,{P}):INDEX(分_贷,{Q}))))')
        D, E = f'{c["dr"]}{r}', f'{c["cr"]}{r}'
        ws[f'{c["ok"]}{r}'] = f'=IF({nil},"",IF(ROUND({D}-{E},2)=0,"√","✗ 差"&FIXED({D}-{E},2)))'
        L = f'{c["lines"]}{r}'
        ws[f'{c["pages"]}{r}'] = f'=IF({nil},"",MAX(1,ROUNDUP({L}/{LINES},0)))'
        pg, p0 = f'{c["pages"]}{r}', f'{c["page0"]}{r}'
        ws[p0] = f'=IF({nil},"",1)' if k == 1 else f'=IF({nil},"",{c["page0"]}{r - 1}+{c["pages"]}{r - 1})'
        ws[f'{c["pk"]}{r}'] = f'=IF({nil},1000000,{p0})'
        ws[f'{c["a4"]}{r}'] = (f'=IF({nil},"",ROUNDUP({p0}/2,0)&IF(ROUNDUP(({p0}+{pg}-1)/2,0)>ROUNDUP({p0}/2,0),'
                               f'"～"&ROUNDUP(({p0}+{pg}-1)/2,0),""))')
        ws[f'{c["src"]}{r}'] = f'=IF({nil},"",IF({N}>10,"月末·","")&INDEX({TB_NAME},{t}))'
        ws[f'{c["memo"]}{r}'] = f'=IF({nil},"",{memo_expr(N, O)})'
    cols = [c[k] for k in ('no', 'date', 'memo', 'dr', 'cr', 'ok', 'lines', 'pages', 'page0', 'a4', 'src')]
    style_rows(ws, VL_R0, VL_R1, cols, auto=tuple(cols),
               fmts={c['date']: DATE, c['dr']: MONEY, c['cr']: MONEY, c['no']: '"记-"000', c['lines']: '0',
                     c['pages']: '0', c['page0']: '0'},
               aligns={c['memo']: AL, c['dr']: AR, c['cr']: AR})
    for r in range(VL_R0, VL_R1 + 1):
        for k in ('key', 'sno', 'n', 'st', 'en', 'pk'):
            ws[f'{c[k]}{r}'].font = F_HELP
    hide(ws, *[c[k] for k in ('key', 'sno', 'n', 'st', 'en', 'pk')], 'T', 'U')
    ws.conditional_formatting.add(f'{c["ok"]}{VL_R0}:{c["ok"]}{VL_R1}',
                                  FormulaRule(formula=[f'LEFT(${c["ok"]}{VL_R0},1)="✗"'], fill=cf_bg('FFFFC7CE'),
                                              font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'A{VL_R0}:{VL_LAST}{VL_R1}',
                                  FormulaRule(formula=[f'AND(ISNUMBER($N{VL_R0}),$N{VL_R0}>10)'],
                                              font=Font(name=YH, sz=10, color='FF7030A0')))
    for nm, k in (('起页键', 'pk'), ('起始页', 'page0'), ('页数', 'pages'), ('来源号', 'sno'), ('n', 'n'), ('起点', 'st'),
                  ('终点', 'en'), ('日期', 'date'), ('摘要', 'memo'), ('借', 'dr'), ('贷', 'cr'), ('条数', 'lines')):
        _dn(wb, f'汇总_{nm}', vrng(k))
    _dn(wb, '汇总_总页数', f"{SH_VLIST}!{S('总页数')}")
    _dn(wb, '汇总_选月', f"{SH_VLIST}!{S('选月')}")
    home_link(ws, 'L1')
    ws.column_dimensions['L'].width = 10
    ws.freeze_panes = f'C{VL_R0}'
    ws.auto_filter.ref = f'A{VL_HDR}:{VL_LAST}{VL_R1}'
    print_setup(ws, f'{VL_HDR}:{VL_HDR}', landscape=False)
    ws.print_area = f'A1:{VL_LAST}{VL_R1}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'


# ═══════════════════════════════ 记账凭证（打印） ═══════════════════════════════
VP = dict(page='I', no='J', word='K', k='L', q='M', src='N', n='O', st='P', en='Q', pp='R', idx='L')


def detail_expr(S, n, L):
    """明细（往来单位、账户、物料、产品…）：S＝来源格、n＝n 格、L＝分_ 第几行格"""
    g = f'INDEX(分_腿,{L})'
    zi = (f'CHOOSE({g},INDEX(账户_名称,MAX(1,INDEX(资_账户号,{n}))),'
          f'IF(INDEX(资_单位,{n})<>"",INDEX(资_单位,{n}),IF(INDEX(资_报销人号,{n})>0,'
          f'INDEX(人员_姓名,INDEX(资_报销人号,{n}))&IF(INDEX(资_报销项目号,{n})>0,"·"&INDEX(报销_名称,INDEX(资_报销项目号,{n})),""),'
          f'"")),"")')
    lk, pk = f'INDEX(进_料号,{n})', f'INDEX(进_品号,{n})'
    jin = (f'CHOOSE({g},IF({lk}>0,INDEX(料档_名称,{lk}),IF({pk}>0,INDEX(品档_名称,{pk}),"")),"",INDEX(进_供应商,{n}),'
           f'INDEX(进_物流单位,{n}))')
    sp, sl = f'INDEX(销_品号,{n})', f'INDEX(销_料号,{n})'
    xs = (f'CHOOSE({g},INDEX(销_客户,{n}),IF({sp}>0,INDEX(品档_名称,{sp}),IF({sl}>0,INDEX(料档_名称,{sl}),"")),"","","",'
          f'INDEX(销_物流单位,{n}))')
    mj = f'INDEX(手_明细,{L}-{B4})'
    pp = (f'CHOOSE({S}-10,IF({g}<={NP},{g},0),IF({g}>1,{g}-1,0),IF({g}<={NP},{g},0),MOD({g}-1,{NP})+1,'
          f'IF({g}>3,{g}-3,0),0,0,0,0,0,0,0)')
    me = f'IF({pp}>0,INDEX(品档_名称,MAX(1,{pp})),"")'
    return (f'IF({S}={SRC_CASH},{zi},IF({S}={SRC_PUR},{jin},IF({S}={SRC_SALE},{xs},IF({S}={SRC_MJ},{mj},{me}))))&""')


def build_vprint(wb, ws):
    c = VP
    widths(ws, {'A': 1.5, 'B': 27, 'C': 12, 'D': 21, 'E': 16, 'F': 16, 'G': 5, 'H': 2, c['page']: 9, c['no']: 8,
                c['word']: 22, 'S': 95})
    ws.sheet_properties.tabColor = C_ACC[2:]
    ws.sheet_view.showGridLines = False
    thin_b = Side(style='thin', color='FF000000')
    YHf = lambda sz, b=False, color='FF000000': Font(name=YH, sz=sz, bold=b, color=color)
    F_NO = YHf(10, True, 'FFC00000')
    last_r = N_VPAGE * VP_ROWS
    sp = '$L$2'                                             # 从第几页起（规范后）
    # 黄格：从第几页起
    ws.merge_cells('I2:J2')
    put(ws, 'I2', '从第几页起：', F_KPI_L, fill('FFDDEBF7'), align=AR, border=False)
    put(ws, 'K2', 1, F_SEL, YELLOW, '0', AC)
    ws['L2'] = '=IF(ISNUMBER($K$2),IF($K$2>=1,INT($K$2),1),1)'
    ws['L2'].font = F_HELP
    dv = DataValidation(type='whole', operator='between', formula1='1', formula2='100000', allow_blank=True,
                        showErrorMessage=True, errorStyle='warning', errorTitle='从第几页起', error='填 1、201、401 这样的页码',
                        showInputMessage=True, promptTitle='提示',
                        prompt=f'这张表一次放 {N_VPAGE} 页凭证；这个月超过 {N_VPAGE} 页的，印完前 {N_VPAGE} 页把这里改成 {N_VPAGE + 1} 再印')
    ws.add_data_validation(dv)
    dv.add('K2')
    H = {1: 33.75, 2: 21.75, 3: 25.5, 12: 40.5, 13: 25.5}
    for i in range(4, 12):
        H[i] = 40.5
    breaks = []
    for p in range(1, N_VPAGE + 1):
        R = (p - 1) * VP_ROWS + 1
        P0 = f'${c["page"]}${R}'
        X = lambda k: f'${c[k]}${R}'
        nil = f'{P0}=""'
        hid = {
            'k': f'=IF({nil},"",MATCH({P0},汇总_起页键,1))',
            'q': f'=IF({nil},"",{P0}-INDEX(汇总_起始页,{X("k")})+1)',
            'src': f'=IF({nil},"",INDEX(汇总_来源号,{X("k")}))',
            'n': f'=IF({nil},"",INDEX(汇总_n,{X("k")}))',
            'st': f'=IF({nil},"",INDEX(汇总_起点,{X("k")}))',
            'en': f'=IF({nil},"",INDEX(汇总_终点,{X("k")}))',
            'pp': f'=IF({nil},"",INDEX(汇总_页数,{X("k")}))',
        }
        ws[f'{c["page"]}{R}'] = f'=IF({sp}+{p - 1}>汇总_总页数,"",{sp}+{p - 1})'
        for k, fm in hid.items():
            ws[f'{c[k]}{R}'] = fm
            ws[f'{c[k]}{R}'].font = F_HELP
        Kk, Q, PP = X('k'), X('q'), X('pp')
        ws[f'{c["no"]}{R}'] = f'=IF({nil},"",{Kk})'
        ws[f'{c["no"]}{R}'].font = F_NO
        ws[f'{c["word"]}{R}'] = f'=IF({nil},"","记-"&TEXT({Kk},"000")&IF({PP}>1,"（"&{Q}&"/"&{PP}&"）",""))'
        ws[f'{c["word"]}{R}'].font = F_NOTE
        ws[f'{c["page"]}{R}'].font = F_NOTE
        # 第 1 行：标题（常量；空页条件格式字变白）
        ws.merge_cells(f'B{R}:G{R}')
        put(ws, f'B{R}', '记  账  凭  证', YHf(18, True), align=AC, border=False)
        # 第 2 行：单位、日期、字号、附单据
        r = R + 1
        put(ws, f'B{r}', f'=IF({nil},"","单位："&P_公司名称)', YHf(9), align=Alignment(horizontal='left', vertical='center',
                                                                                         shrink_to_fit=True), border=False)
        ws.merge_cells(f'C{r}:D{r}')
        put(ws, f'C{r}', f'=IF({nil},"",INDEX(汇总_日期,{Kk}))', YHf(11, True), None, 'yyyy"年"m"月"d"日"', AC, border=False)
        ws.merge_cells(f'E{r}:F{r}')
        put(ws, f'E{r}', (f'=IF({nil},"","记 字第 "&IF(P_凭证号打印=1,TEXT({Kk},"000"),"　　　　")&" 号"'
                          f'&IF({PP}>1,"（共"&{PP}&"页 第"&{Q}&"页）",""))'), YHf(11, True), align=AC, border=False)
        put(ws, f'G{r}', '附　 张', YHf(9), align=AC, border=False)
        # 第 3 行：表头（常量）
        r = R + 2
        for col, t in (('B', '摘  要'), ('C', '总账科目'), ('D', '明细科目'), ('E', '借方金额'), ('F', '贷方金额'), ('G', '√')):
            put(ws, f'{col}{r}', t, YHf(11, True), align=AC, border=False)
        # 8 条分录
        for i in range(1, LINES + 1):
            r = R + 2 + i
            Lr = f'${c["idx"]}{r}'
            ws[f'{c["idx"]}{r}'] = (f'=IF({nil},"",IFERROR(IF({X("src")}={SRC_MJ},{B4}+MATCH({X("n")}*1000+({Q}-1)*{LINES}+{i},'
                                    f'手_印键,0),{X("st")}-1+MATCH(({Q}-1)*{LINES}+{i},INDEX(分_累,{X("st")}):INDEX(分_累,{X("en")}),0)),""))')
            ws[f'{c["idx"]}{r}'].font = F_HELP
            no = f'INDEX(分_科目号,{Lr})'
            emp = f'{Lr}=""'
            put(ws, f'B{r}', (f'=IF({emp},"",IF({X("src")}={SRC_MJ},INDEX({inref(SH_MJ, MJ_COLS["摘要"])},{Lr}-{B4}+1)&"",'
                              f'INDEX(汇总_摘要,{Kk})))'), YHf(10), align=ALW, border=False)
            put(ws, f'C{r}', f'=IF({emp},"",IF(INDEX(科目_一级号,MAX(1,{no}))=0,"",INDEX(科目_名称,INDEX(科目_一级号,MAX(1,{no})))))',
                YHf(10.5), align=ALW, border=False)
            n2 = f'IF(INDEX(科目_级次,MAX(1,{no}))=2,"-"&INDEX(科目_名称,MAX(1,{no})),"")'
            det = detail_expr(X('src'), X('n'), Lr)
            put(ws, f'D{r}', f'=IF({emp},"",MID({n2}&IF({det}<>"","-"&{det},""),2,200))', YHf(10.5), align=ALW, border=False)
            for col, side in (('E', '借'), ('F', '贷')):
                v = f'INDEX(分_{side},{Lr})'
                put(ws, f'{col}{r}', f'=IF({emp},"",IF({v}=0,"",{v}))', YHf(11), None, MONEY,
                    Alignment(horizontal='right', vertical='center', shrink_to_fit=True), border=False)
            put(ws, f'G{r}', None, YHf(11), border=False)
        # 合计
        r = R + 3 + LINES
        last = f'{Q}={PP}'
        put(ws, f'B{r}', f'=IF({nil},"",IF({last},"合    计","（接下页）"))', YHf(11, True), align=AC, border=False)
        ws.merge_cells(f'C{r}:D{r}')
        put(ws, f'C{r}', None, YHf(11, True), border=False)
        for col, nm in (('E', '汇总_借'), ('F', '汇总_贷')):
            put(ws, f'{col}{r}', f'=IF({nil},"",IF({last},INDEX({nm},{Kk}),""))', YHf(11, True), None, MONEY,
                Alignment(horizontal='right', vertical='center', shrink_to_fit=True), border=False)
        put(ws, f'G{r}', None, YHf(11, True), border=False)
        # 签字（一格）
        r = R + 4 + LINES
        ws.merge_cells(f'B{r}:G{r}')
        put(ws, f'B{r}', (f'=IF({nil},"","会计主管："&P_会计主管&"　　　　记账："&P_记账&"　　　　出纳："&P_出纳'
                          f'&"　　　　审核："&P_审核&"　　　　制单："&P_制单)'), YHf(10), align=AL, border=False)
        for i in range(1, VP_ROWS):
            ws.row_dimensions[R + i - 1].height = H[i]
        ws.row_dimensions[R + VP_ROWS - 1].height = 60 if p % 2 == 1 else 6
    # 条件格式：有凭证的页才有框；空页的常量字变白
    rel = f'MOD(ROW()-1,{VP_ROWS})'
    used = f'INDEX(${c["page"]}$1:${c["page"]}${last_r},ROW()-{rel})<>""'
    box = Border(left=thin_b, right=thin_b, top=thin_b, bottom=thin_b)
    rng = f'B1:G{last_r}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({used},{rel}>=2,{rel}<={2 + LINES})'], border=box))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({used},{rel}={3 + LINES})'], border=box))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({used},{rel}=2)'], fill=cf_bg('FFF2F2F2')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'NOT({used})'], font=Font(color='FFFFFFFF')))
    ws.row_breaks = RowBreak()
    for p in range(1, N_VPAGE // 2 + (N_VPAGE % 2)):
        ws.row_breaks.append(Break(id=2 * VP_ROWS * p))
    ps = ws.page_setup
    ps.paperSize = 9
    ps.orientation = 'portrait'
    ps.scale = 75
    ps.fitToWidth = ps.fitToHeight = None
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    pm = ws.page_margins
    pm.left, pm.right, pm.top, pm.bottom, pm.header, pm.footer = 0.6, 0.3, 0.4, 0.3, 0.2, 0.2
    ws.print_options.horizontalCentered = True
    ws.print_area = f'B1:G{last_r}'
    ws._keep_print = True                                   # make.py 别再套横向一页宽的打印设置
    tips = ['★ 一张 A4 纸上下印两页凭证（中间留宽，裁开分别装订）。凭证跟着【凭证汇总】选的月份走；打印页码范围填 1 到【凭证汇总】右上角「A4 张数」。',
            f'I＝凭证第几页，J＝凭证号（红字，可照着手写），K＝字号。这张表一次放 {N_VPAGE} 页；这个月超过 {N_VPAGE} 页的，'
            f'印完前 {N_VPAGE} 页把 K2「从第几页起」改成 {N_VPAGE + 1} 再印。',
            '纸张已设好：A4 纵向、缩放 75%、左边留装订边。别改成「调整为一页」，不然两张凭证会挤在一起。一张凭证超过 8 条分录接着印下一页，两页是同一个号。',
            '凭证号印不印、签字栏的人名在【基础资料】参数里改；附单据张数手写。']
    for i, t in enumerate(tips):
        ws[f'S{4 + i}'] = t
        ws[f'S{4 + i}'].font = F_TIP
        ws[f'S{4 + i}'].alignment = ALW
    hide(ws, 'L', 'M', 'N', 'O', 'P', 'Q', 'R')
    ws.freeze_panes = None


# ═══════════════════════════════ 科目余额表 ═══════════════════════════════
TB = dict(code='A', name='B', dir='C', lvl='D', ob_d='E', ob_c='F', m_d='G', m_c='H', y_d='I', y_c='J', cb_d='K', cb_c='L',
          no='N', onet='O', cnet='P')
TB_LAST = 'L'
TB_HDR, TB_R0 = 5, 6
TB_R1 = TB_R0 + N_COA - 1
TB_SC = ['输入', '认出', '选月', '漏借', '漏贷', '期初差']


def tsc(k):
    return f'$U${3 + TB_SC.index(k)}'


def build_tb(wb, ws):
    t = TB
    widths(ws, {t['code']: 10, t['name']: 30, t['dir']: 5, t['lvl']: 5, **{t[k]: 15 for k in
                                                                          ('ob_d', 'ob_c', 'm_d', 'm_c', 'y_d', 'y_c', 'cb_d', 'cb_c')}})
    title(ws, '科 目 余 额 表', TB_LAST, C_ACC,
          '★ 选一个月（B3 黄格，空着＝截止日期那个月）：每个科目的月初余额、本月发生、本年累计发生（1 月到这个月）、月末余额，全部从记账凭证汇总；'
          '一级科目＝下级合计＋直接记在一级上的。余额按借贷拆开列（借方余额列借方，贷方余额列贷方）。损益类科目不结转（表结法），本年累计就是利润表的数；'
          '未分配利润在资产负债表里倒挤。最下面一行是一级科目合计，借贷应该相等（F3 看平不平）。只能筛选、不能排序。')
    ws.row_dimensions[3].height = 26
    for i in range(1, CI(TB_LAST) + 1):
        put(ws, f'{CL(i)}3', None, fill_=fill('FFDDEBF7'), border=False)
    selector(ws, 'A3', '月份', 'B3', None, '"1,2,3,4,5,6,7,8,9,10,11,12"',
             prompt='选 1～12（本年度）；空着＝截止日期那个月；填 2026-9、9月 也认')
    S = tsc
    rows_lvl1 = lambda col: f'SUMIF(${t["lvl"]}${TB_R0}:${t["lvl"]}${TB_R1},1,{col}{TB_R0}:{col}{TB_R1})'
    f = {
        '输入': '=' + month_raw('$B$3'),
        '认出': f'=IF({S("输入")}=-1,1,IF(AND({S("输入")}>=1,{S("输入")}<=12,INT({S("输入")})={S("输入")}),1,0))',
        '选月': f'=IF(OR({S("输入")}=-1,{S("认出")}=0),P_截止月,{S("输入")})',
        '漏借': f'=ROUND(SUM(分_借)-SUMIF(科目_有效,1,矩_全年借),2)',
        '漏贷': f'=ROUND(SUM(分_贷)-SUMIF(科目_有效,1,矩_全年贷),2)',
        '期初差': f'=ROUND(SUM(矩_期初),2)',
    }
    for i, k in enumerate(TB_SC):
        _lab(ws, f'T{3 + i}', k)
        ws[f'U{3 + i}'] = f[k]
        ws[f'U{3 + i}'].font = F_HELP
    ws.merge_cells('C3:D3')
    put(ws, 'C3', (f'="实际："&P_年度&" 年 "&{S("选月")}&" 月"&IF({S("认出")}=0,"　⚠ 月份没认出来，按截止日期那个月","")'),
        F_NOTE, align=ALW, border=False)
    put(ws, 'E3', '平衡检查：', F_KPI_L, fill('FFDDEBF7'), align=AR, border=False)
    tr = TB_R1 + 1
    pair = lambda a, b: f'ROUND({a}{tr}-{b}{tr},2)=0'
    chk = (f'=IF(AND({pair(t["ob_d"], t["ob_c"])},{pair(t["m_d"], t["m_c"])},{pair(t["y_d"], t["y_c"])},'
           f'{pair(t["cb_d"], t["cb_c"])},{S("漏借")}=0,{S("漏贷")}=0),"√ 借贷平衡（月初、本月、本年、月末）",'
           f'IF(OR({S("漏借")}<>0,{S("漏贷")}<>0),"✗ 有分录记到了科目表里没有（或有问题）的科目：借 "&FIXED({S("漏借")},2)&"，贷 "&FIXED({S("漏贷")},2),'
           f'IF({S("期初差")}<>0,"✗ 期初不平：差 "&FIXED({S("期初差")},2)&"（看【基础资料】⑧ 和各期初）",'
           f'"✗ 借贷不平：月末差 "&FIXED({t["cb_d"]}{tr}-{t["cb_c"]}{tr},2))))')
    ws.merge_cells('F3:L3')
    put(ws, 'F3', chk, Font(name=YH, sz=12, bold=True, color='FFC00000'), fill('FFDDEBF7'), align=AL, border=False)
    ws.conditional_formatting.add('F3', FormulaRule(formula=['LEFT($F$3,1)="√"'], font=Font(name=YH, sz=12, bold=True,
                                                                                               color='FF548235')))
    for c0, c1, txt_ in ((t['ob_d'], t['ob_c'], '月初余额'), (t['m_d'], t['m_c'], '本月发生额'),
                         (t['y_d'], t['y_c'], '本年累计发生额'), (t['cb_d'], t['cb_c'], '月末余额')):
        ws.merge_cells(f'{c0}{TB_HDR - 1}:{c1}{TB_HDR - 1}')
        put(ws, f'{c0}{TB_HDR - 1}', txt_, F_HDR, fill(C_ACC), align=AC)
        put(ws, f'{c1}{TB_HDR - 1}', None, fill_=fill(C_ACC))
    for col in (t['code'], t['name'], t['dir'], t['lvl']):
        put(ws, f'{col}{TB_HDR - 1}', None, fill_=fill(C_ACC))
    header(ws, TB_HDR, [(t['code'], '科目编码'), (t['name'], '科目名称'), (t['dir'], '方向'), (t['lvl'], '级次'),
                        (t['ob_d'], '借方'), (t['ob_c'], '贷方'), (t['m_d'], '借方'), (t['m_c'], '贷方'), (t['y_d'], '借方'),
                        (t['y_c'], '贷方'), (t['cb_d'], '借方'), (t['cb_c'], '贷方')], C_ACC, height=24)
    for k, txt_ in (('no', '科目号'), ('onet', '月初净'), ('cnet', '月末净')):
        _lab(ws, f'{t[k]}{TB_HDR}', txt_)
    s = S('选月')
    for k in range(1, N_COA + 1):
        r = TB_R0 + k - 1
        M = f'${t["no"]}{r}'
        nil = f'{M}=""'
        ws[M[1:]] = f'=IFERROR(MOD(SMALL(矩_排序键,ROW()-{TB_HDR}),1000),"")'
        ws[f'{t["code"]}{r}'] = f'=IF({nil},"",INDEX(科目_编码,{M}))'
        ws[f'{t["name"]}{r}'] = f'=IF({nil},"",REPT("　",INDEX(科目_级次,{M})-1)&INDEX(科目_名称,{M}))'
        ws[f'{t["dir"]}{r}'] = f'=IF({nil},"",INDEX(科目_方向,{M}))'
        ws[f'{t["lvl"]}{r}'] = f'=IF({nil},"",INDEX(科目_级次,{M}))'
        cum = lambda nm, a, b: f'SUM(INDEX({nm},{M},{a}):INDEX({nm},{M},{b}))'
        ws[f'{t["onet"]}{r}'] = (f'=IF({nil},"",ROUND(INDEX(矩_汇期初,{M})+IF({s}>1,{cum("矩_汇借", 1, f"{s}-1")}'
                                 f'-{cum("矩_汇贷", 1, f"{s}-1")},0),2))')
        ws[f'{t["m_d"]}{r}'] = f'=IF({nil},"",INDEX(矩_汇借,{M},{s}))'
        ws[f'{t["m_c"]}{r}'] = f'=IF({nil},"",INDEX(矩_汇贷,{M},{s}))'
        ws[f'{t["y_d"]}{r}'] = f'=IF({nil},"",{cum("矩_汇借", 1, s)})'
        ws[f'{t["y_c"]}{r}'] = f'=IF({nil},"",{cum("矩_汇贷", 1, s)})'
        on, cn = f'{t["onet"]}{r}', f'{t["cnet"]}{r}'
        ws[cn] = f'=IF({nil},"",ROUND({on}+{t["m_d"]}{r}-{t["m_c"]}{r},2))'
        ws[f'{t["ob_d"]}{r}'] = f'=IF({nil},"",MAX({on},0))'
        ws[f'{t["ob_c"]}{r}'] = f'=IF({nil},"",MAX(-{on},0))'
        ws[f'{t["cb_d"]}{r}'] = f'=IF({nil},"",MAX({cn},0))'
        ws[f'{t["cb_c"]}{r}'] = f'=IF({nil},"",MAX(-{cn},0))'
    cols = [t[k] for k in ('code', 'name', 'dir', 'lvl', 'ob_d', 'ob_c', 'm_d', 'm_c', 'y_d', 'y_c', 'cb_d', 'cb_c')]
    money = cols[4:]
    style_rows(ws, TB_R0, TB_R1, cols, auto=tuple(cols), fmts={**{x: MONEY for x in money}, t['code']: '@'},
               aligns={t['code']: AL, t['name']: AL, **{x: AR for x in money}})
    for r in range(TB_R0, TB_R1 + 1):
        for k in ('no', 'onet', 'cnet'):
            ws[f'{t[k]}{r}'].font = F_HELP
    hide(ws, t['no'], t['onet'], t['cnet'], 'T', 'U')
    ws.conditional_formatting.add(f'A{TB_R0}:{TB_LAST}{TB_R1}',
                                  FormulaRule(formula=[f'${t["lvl"]}{TB_R0}=1'], font=Font(name=YH, sz=10, bold=True),
                                              fill=cf_bg('FFEEF3FA')))
    r = tr
    put(ws, f'{t["code"]}{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'{t["name"]}{r}', '（一级科目合计，借＝贷）', F_TXTB, FILL_TOT, align=AL)
    for col in (t['dir'], t['lvl']):
        put(ws, f'{col}{r}', None, F_TXTB, FILL_TOT)
    for col in money:
        put(ws, f'{col}{r}', f'={rows_lvl1(col)}', F_TXTB, FILL_TOT, MONEY, AR)
    home_link(ws, 'M1')
    ws.column_dimensions['M'].width = 10
    ws.freeze_panes = f'C{TB_R0}'
    ws.auto_filter.ref = f'A{TB_HDR}:{TB_LAST}{TB_R1}'
    print_setup(ws, f'{TB_HDR - 1}:{TB_HDR}', landscape=True)
    ws.print_area = f'A1:{TB_LAST}{tr}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'


# ═══════════════════════════════ 年末结转 ═══════════════════════════════
YE_LAST = 'S'
HX0 = 'Y'                                  # 辅助区从 Y 列起
N7 = 3 * N_LEND                            # ⑦ 槽位（每个出借人 借款/垫付/投资 各一）
N8A = 60                                   # ⑧ 资金台帐里持有 1801 的付款，最多列 60 笔
N8_OUT = N_OP + 30                         # ⑧ 输出行数
FAM8 = ['1001', '1002', '1012', '1121', '1122', '2202', '1403', '1405', '1601', '1602', '1801', '1123', '2001', '3001', '2231',
        '2232', '4001', '3104']
EXACT8 = ['224101', '224102', '122103']


class Hx:
    """年末结转右边的辅助表：每张表一组列（从 HX0 起往右排），第 2 行起"""

    def __init__(self):
        self.c = CI(HX0)

    def cols(self, names):
        out = {k: CL(self.c + i) for i, k in enumerate(names)}
        self.c += len(names) + 1
        return out


def ye_copy(src_ref, n, guard):
    v = f'INDEX({src_ref},{n}+1)'
    return f'=IF({guard},"",IF({v}="","",{v}))'


YN = 'W'                                   # 年末结转：每块数据行的第几个（常量，公式都引用它 → 整块共享公式）


def build_ye(wb, ws):
    widths(ws, {'A': 16, 'B': 12, 'C': 14, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 13, 'I': 13, 'J': 13, 'K': 13,
                'L': 10, 'M': 10, 'N': 10, 'O': 10, 'P': 10, 'Q': 9, 'R': 11, 'S': 14})
    title(ws, '年 末 结 转（给下一本的期初）', YE_LAST, C_ACC,
          '★ 一本账＝一年。到 12 月底（首页截止日期填 12/31 或空着、12 月的单子都登完），用这张表开下一年的新本：新本【基础资料】建账日期填下一年 1 月 1 日，'
          '再把下面每一块（灰色表头下面的数据行）整块选中 → 复制 → 到新本对应的那一块第一行 → 右键「选择性粘贴」→「数值」。'
          '每块的列顺序跟新本那一块一模一样。① ② 账户期末余额（垫付、备用金账户不填，在 ⑦ ⑧）；② ⑦ 每个出借人一行本金（日期＝新本建账日，已付利息填负数＝欠着的利息）；'
          '③ ⑧ 其他科目余额（长期待摊、预付带剩余月数和摊进，新本接着摊；制造费用留存、备用金按人）；④ 往来单位期末应收应付、没开票、欠票；⑤ ⑥ 物料档案期末数量金额、在产；'
          '⑦ 固定资产（建账前已提＝到年底的累计折旧；今年处置的不带）。损益、利润分配不用带（新本未分配利润倒挤）。每块下面一行跟账上核对。')
    ws.row_dimensions[3].height = 30
    ws.merge_cells(f'A3:{YE_LAST}3')
    put(ws, 'A3', (f'=IF(P_截止月<12,"⚠ 现在截止日期是 "&TEXT(P_截止,"yyyy-mm-dd")&"，12 月的折旧、摊销、利息、工资、成本还没算完："'
                   f'&"下面是到现在的数，正式结转请等 12 月的单子登完（首页截止日期填 "&P_年度&"-12-31 或空着）",'
                   f'"✓ 截止到 "&TEXT(P_截止,"yyyy-mm-dd")&"，可以结转。新本【基础资料】建账日期填 "&(P_年度+1)&"-01-01")'),
        F_RED, fill('FFFFF2CC'), align=ALW, border=False)
    hx = Hx()
    bal = lambda code: f'INDEX(矩_汇年末,{K(code)})'
    r = 5
    chk_cells = []

    def sec(r, txt_, cols_map, last_col, color=C_ACC):
        ws.merge_cells(f'A{r}:{last_col}{r}')
        put(ws, f'A{r}', txt_, F_SEC, fill(color), align=AL)
        for i in range(2, CI(last_col) + 1):
            ws.cell(row=r, column=i).border = BD
        header(ws, r + 1, [(c, k) for k, c in cols_map.items()], 'FF7F7F7F', height=30)
        return r + 2

    def check_row(r, last_col, parts):
        """parts: [(说明, 表里合计, 账上)]"""
        ws.merge_cells(f'A{r}:{last_col}{r}')
        segs = []
        for lab, a, b in parts:
            d = f'ROUND(({a})-({b}),2)'
            segs.append(f'"{lab}："&IF({d}=0,"√ 跟账上一致（"&FIXED({b},2)&"）","✗ 表里 "&FIXED({a},2)&"，账上 "&FIXED({b},2)&"，差 "&FIXED({d},2))')
        put(ws, f'A{r}', '=' + '&"　　"&'.join(segs), F_NOTE, fill('FFF2F2F2'), align=ALW, border=False)
        ws.row_dimensions[r].height = 30
        chk_cells.append(f'A{r}')
        return r + 2

    def money_cells(r0, r1, cols):
        for rr in range(r0, r1 + 1):
            for col in cols:
                ws[f'{col}{rr}'].number_format = MONEY

    # ── ① ② 账户 ──
    r1 = sec(r, '① 【基础资料】② 账户：期末余额（填进新本 ② 的「期初余额」；垫付、备用金账户不填）', ACC_COLS, 'E')
    acc = lambda k: inref(SH_BASE, ACC_COLS[k], block='账户')
    for i in range(1, N_ACC + 1):
        rr = r1 + i - 1
        ws[f'{YN}{rr}'] = i
        n_ = f'${YN}{rr}'
        g = f'INDEX(账户_名称,{n_})=""'
        ws[f'A{rr}'] = f'=IF({g},"",INDEX(账户_名称,{n_}))'
        for k in ('类型', '人员', '备注'):
            ws[f'{ACC_COLS[k]}{rr}'] = ye_copy(acc(k), n_, g)
        ty = f'INDEX(账户_类型号,{n_})'
        ws[f'C{rr}'] = (f'=IF(OR({g},{ty}<1,{ty}>4),"",ROUND(INDEX(账户_期初,{n_})+SUMIFS(资_净额,资_账户号,{n_},资_有效,1),2))')
    style_rows(ws, r1, r1 + N_ACC - 1, list('ABCDE'), auto=tuple('ABCDE'), fmts={'C': MONEY})
    rr = r1 + N_ACC
    cash_bal = '+'.join(bal(c) for c in ('1001', '1002', '1012', '1121'))
    r = check_row(rr, 'E', [('现金、银行、微信、承兑合计', f'SUM(C{r1}:C{rr - 1})', cash_bal)])

    # ── ② ⑦ 出借人 ──
    L = hx.cols(['l', '名称', '单位号', '投本', '垫本', '借本', '借欠息', '投欠息'])
    Lr = lambda k: f'${L[k]}$2:${L[k]}${N_LEND + 1}'
    for k, col in L.items():
        _lab(ws, f'{col}1', '⑦' + k)
    for l in range(1, N_LEND + 1):
        w = l + 1
        A = f'${L["l"]}{w}'
        ws[A[1:]] = l
        ws[f'{L["名称"]}{w}'] = f'=INDEX(出借_名称,{A})&""'
        u = f'{L["单位号"]}{w}'
        ws[u] = f'=N(INDEX(出借_单位号,{A}))'
        lo = lambda kind: f'SUMIFS(借期_金额,借期_出借人号,{A},借期_性质,{kind},借期_有效,1)'
        ws[f'{L["投本"]}{w}'] = f'=ROUND({lo(2)}+SUMIFS(资_本金变动,资_出借人号,{A},资_性质,2),2)'
        k224101 = K('224101')
        ws[f'{L["垫本"]}{w}'] = (f'=ROUND({lo(3)}+SUMIFS(资_本金变动,资_出借人号,{A},资_账户科目号,{k224101})'
                               f'+SUMIFS(资_本金变动,资_出借人号,{A},资_账户科目号,"<>"&{k224101},资_对方号,{k224101}),2)')
        ws[f'{L["借本"]}{w}'] = (f'=ROUND({lo(chr(34) + "<>2" + chr(34))}+SUMIFS(资_本金变动,资_出借人号,{A},资_性质,1)'
                               f'-{L["垫本"]}{w},2)')
        for nm, pre, code in (('借欠息', '借', '2231'), ('投欠息', '投', '2232')):
            ws[f'{L[nm]}{w}'] = (f'=ROUND(INDEX(息_{pre}期初应付息,{A})+INDEX(息_{pre}累计利息,{A})'
                                f'+IF({u}=0,0,SUMIFS(资_净额,资_对方号,{K(code)},资_单位号,{u})),2)')
    Q = hx.cols(['s', '金额', '已付', '出', '累'])
    for k, col in Q.items():
        _lab(ws, f'{col}1', '⑦槽' + k)
    for s in range(1, N7 + 1):
        w = s + 1
        l, kind = (s - 1) // 3 + 1, (s - 1) % 3 + 1
        lw = l + 1
        a, d, b = f'${L["借本"]}${lw}', f'${L["垫本"]}${lw}', f'${L["投本"]}${lw}'
        i1, i2 = f'${L["借欠息"]}${lw}', f'${L["投欠息"]}${lw}'
        ws[f'{Q["s"]}{w}'] = s
        amt, paid, out = {
            1: (a, f'IF(OR({a}<>0,{d}=0),-{i1},0)', f'OR({a}<>0,AND({d}=0,{i1}<>0))'),
            2: (d, f'IF({a}=0,-{i1},0)', f'{d}<>0'),
            3: (b, f'-{i2}', f'OR({b}<>0,{i2}<>0)'),
        }[kind]
        ws[f'{Q["金额"]}{w}'] = f'={amt}'
        ws[f'{Q["已付"]}{w}'] = f'={paid}'
        ws[f'{Q["出"]}{w}'] = f'=IF({out},1,0)'
        ws[f'{Q["累"]}{w}'] = f'={Q["出"]}{w}' + (f'+{Q["累"]}{w - 1}' if s > 1 else '')
    qr = lambda k: f'${Q[k]}$2:${Q[k]}${N7 + 1}'
    r1 = sec(r, '② 【基础资料】⑦ 建账前的借款、投资、垫付：每个出借人一行本金（日期＝新本建账日；已付利息是负数＝到年底还欠的利息）', LO_COLS, 'F')
    for j in range(1, N7 + 1):
        rr = r1 + j - 1
        ws[f'{YN}{rr}'] = j
        sl = f'IFERROR(MATCH(${YN}{rr},{qr("累")},0),0)'
        ws[f'X{rr}'] = f'={sl}'
        ws[f'X{rr}'].font = F_HELP
        S_ = f'$X{rr}'
        g = f'{S_}=0'
        ws[f'A{rr}'] = f'=IF({g},"",INDEX({Lr("名称")},INT(({S_}-1)/3)+1))'
        ws[f'B{rr}'] = f'=IF({g},"",DATE(P_年度+1,1,1))'
        ws[f'C{rr}'] = f'=IF({g},"",INDEX({qr("金额")},{S_}))'
        ws[f'D{rr}'] = f'=IF({g},"",CHOOSE(MOD({S_}-1,3)+1,"借款","垫付","投资"))'
        ws[f'E{rr}'] = f'=IF({g},"",IF(INDEX({qr("已付")},{S_})=0,"",INDEX({qr("已付")},{S_})))'
        ws[f'F{rr}'] = f'=IF({g},"","上年结转")'
    style_rows(ws, r1, r1 + N7 - 1, list('ABCDEF'), auto=tuple('ABCDEF'), fmts={'B': DATE, 'C': MONEY, 'E': MONEY})
    rr = r1 + N7
    kind_sum = lambda kd: f'SUMIF(D{r1}:D{rr - 1},"{kd}",C{r1}:C{rr - 1})'
    paid = lambda kinds: '+'.join(f'SUMIF(D{r1}:D{rr - 1},"{kd}",E{r1}:E{rr - 1})' for kd in kinds)
    r = check_row(rr, 'F', [('借款', kind_sum('借款'), f'-({bal("2001")}+{bal("224102")})'),
                            ('垫付', kind_sum('垫付'), f'-{bal("224101")}'), ('投资', kind_sum('投资'), f'-{bal("3001")}'),
                            ('欠利息', paid(['借款', '垫付']), bal('2231')), ('欠投资利息', paid(['投资']), bal('2232'))])

    # ── ③ ⑧ 其他期初余额 ──
    E = 'P_截止月'
    A8 = hx.cols(['j', '行', '始', '止', '月摊', '尾', '目标', '剩月', '剩额'])
    for k, col in A8.items():
        _lab(ws, f'{col}1', '1801付' + k)
    for j in range(1, N8A + 1):
        w = j + 1
        ws[f'{A8["j"]}{w}'] = j
        pos = f'{A8["行"]}{w}'
        if j == 1:
            ws[pos] = '=IFERROR(MATCH(1,资_分摊,0),0)'
        else:
            pv = f'{A8["行"]}{w - 1}'
            ws[pos] = f'=IF({pv}=0,0,IFERROR({pv}+MATCH(1,INDEX(资_分摊,{pv}+1):INDEX(资_分摊,{N_CASH}),0),0))'
        for k, nm in (('始', '资_摊始月'), ('止', '资_摊止月'), ('月摊', '资_月摊额'), ('尾', '资_摊尾差'), ('目标', '资_摊目标号')):
            ws[f'{A8[k]}{w}'] = f'=IF({pos}=0,0,INDEX({nm},{pos}))'
        a_, b_ = f'{A8["始"]}{w}', f'{A8["止"]}{w}'
        rm = f'{A8["剩月"]}{w}'
        ws[rm] = f'=IF({pos}=0,0,{b_}-{a_}+1-MAX(0,MIN({E},{b_})-{a_}+1))'
        ws[f'{A8["剩额"]}{w}'] = f'=IF({rm}<=0,0,ROUND({A8["月摊"]}{w}*{rm}+IF({b_}>{E},{A8["尾"]}{w},0),2))'
    B8 = hx.cols(['n', '额', '月数', '始', '止', '月摊', '尾', '目标', '剩月', '剩额'])
    for k, col in B8.items():
        _lab(ws, f'{col}1', '⑧1801' + k)
    for n in range(1, N_OP + 1):
        w = n + 1
        ws[f'{B8["n"]}{w}'] = n
        x = lambda f_: f'INDEX(期初_{f_},{n})'
        amt = f'{B8["额"]}{w}'
        ws[amt] = (f'=IF(AND({x("有效")}=1,{x("科目号")}>0),IF(INDEX(科目_一级号,MAX(1,{x("科目号")}))={K("1801")},'
                   f'ROUND({x("借")}-{x("贷")},2),0),0)')
        ws[f'{B8["月数"]}{w}'] = f'=IF({x("剩余月数")}>0,{x("剩余月数")},P_摊销月数)'
        ws[f'{B8["始"]}{w}'] = '=MAX(P_建账月,P_投产月)'
        ws[f'{B8["止"]}{w}'] = f'={B8["始"]}{w}+{B8["月数"]}{w}-1'
        ws[f'{B8["月摊"]}{w}'] = f'=IF({amt}=0,0,ROUND({amt}/{B8["月数"]}{w},2))'
        ws[f'{B8["尾"]}{w}'] = f'=IF({amt}=0,0,ROUND({amt}-{B8["月摊"]}{w}*{B8["月数"]}{w},2))'
        ws[f'{B8["目标"]}{w}'] = f'=IF({x("摊进")}=2,{K("560207")},{K("410106")})'
        a_, b_ = f'{B8["始"]}{w}', f'{B8["止"]}{w}'
        rm = f'{B8["剩月"]}{w}'
        ws[rm] = f'=IF({amt}=0,0,{b_}-{a_}+1-MAX(0,MIN({E},{b_})-{a_}+1))'
        ws[f'{B8["剩额"]}{w}'] = f'=IF({rm}<=0,0,ROUND({B8["月摊"]}{w}*{rm}+IF({b_}>{E},{B8["尾"]}{w},0),2))'
    # 预付 1123：资金台帐里 分摊＝2 的付款（按月摊）、⑧ 期初 1123 行；按截止月算还剩几个月、还剩多少（跟 1801 一样）
    A8b = hx.cols(['j', '行', '始', '止', '月摊', '尾', '目标', '剩月', '剩额'])
    for k, col in A8b.items():
        _lab(ws, f'{col}1', '1123付' + k)
    for j in range(1, N8A + 1):
        w = j + 1
        ws[f'{A8b["j"]}{w}'] = j
        pos = f'{A8b["行"]}{w}'
        if j == 1:
            ws[pos] = '=IFERROR(MATCH(2,资_分摊,0),0)'
        else:
            pv = f'{A8b["行"]}{w - 1}'
            ws[pos] = f'=IF({pv}=0,0,IFERROR({pv}+MATCH(2,INDEX(资_分摊,{pv}+1):INDEX(资_分摊,{N_CASH}),0),0))'
        for k, nm in (('始', '资_摊始月'), ('止', '资_摊止月'), ('月摊', '资_月摊额'), ('尾', '资_摊尾差'), ('目标', '资_摊目标号')):
            ws[f'{A8b[k]}{w}'] = f'=IF({pos}=0,0,INDEX({nm},{pos}))'
        a_, b_ = f'{A8b["始"]}{w}', f'{A8b["止"]}{w}'
        rm = f'{A8b["剩月"]}{w}'
        ws[rm] = f'=IF({pos}=0,0,{b_}-{a_}+1-MAX(0,MIN({E},{b_})-{a_}+1))'
        ws[f'{A8b["剩额"]}{w}'] = f'=IF({rm}<=0,0,ROUND({A8b["月摊"]}{w}*{rm}+IF({b_}>{E},{A8b["尾"]}{w},0),2))'
    B8b = hx.cols(['n', '额', '可', '月数', '始', '止', '月摊', '尾', '剩月', '剩额', '摊进'])
    for k, col in B8b.items():
        _lab(ws, f'{col}1', '⑧1123' + k)
    op_tj = inref(SH_BASE, OP_COLS['摊进'], block='期初')
    for n in range(1, N_OP + 1):
        w = n + 1
        ws[f'{B8b["n"]}{w}'] = n
        x = lambda f_: f'INDEX(期初_{f_},{n})'
        amt, ok = f'{B8b["额"]}{w}', f'{B8b["可"]}{w}'
        ws[amt] = (f'=IF(AND({x("有效")}=1,{x("科目号")}>0),IF(INDEX(科目_一级号,MAX(1,{x("科目号")}))={K("1123")},'
                   f'ROUND({x("借")}-{x("贷")},2),0),0)')
        ws[ok] = f'=IF(AND({amt}<>0,{x("摊进号")}>0,{x("剩余月数")}>0),1,0)'
        mon = f'{B8b["月数"]}{w}'
        ws[mon] = f'={x("剩余月数")}'
        ws[f'{B8b["始"]}{w}'] = '=P_建账月'
        ws[f'{B8b["止"]}{w}'] = f'={B8b["始"]}{w}+{mon}-1'
        ws[f'{B8b["月摊"]}{w}'] = f'=IF({ok}=0,0,ROUND({amt}/{mon},2))'
        ws[f'{B8b["尾"]}{w}'] = f'=IF({ok}=0,0,ROUND({amt}-{B8b["月摊"]}{w}*{mon},2))'
        a_, b_ = f'{B8b["始"]}{w}', f'{B8b["止"]}{w}'
        rm = f'{B8b["剩月"]}{w}'
        ws[rm] = f'=IF({ok}=1,{b_}-{a_}+1-MAX(0,MIN({E},{b_})-{a_}+1),IF({amt}=0,0,{mon}))'
        ws[f'{B8b["剩额"]}{w}'] = (f'=IF({ok}=1,IF({rm}<=0,0,ROUND({B8b["月摊"]}{w}*{rm}+IF({b_}>{E},{B8b["尾"]}{w},0),2)),'
                                 f'{amt})')
        tj = x('摊进号')
        ws[f'{B8b["摊进"]}{w}'] = (f'=IF({ok}=1,INDEX(科目_编码,{tj})&" "&INDEX(科目_名称,{tj}),'
                                 f'TRIM(INDEX({op_tj},{n}+1)&""))')
    # 备用金账户（122103 按人）
    C8 = hx.cols(['a', '类型', '人', '首', '余额'])
    for k, col in C8.items():
        _lab(ws, f'{col}1', '备用金' + k)
    k122103 = K('122103')
    for a in range(1, N_ACC + 1):
        w = a + 1
        ws[f'{C8["a"]}{w}'] = a
        ty, per = f'{C8["类型"]}{w}', f'{C8["人"]}{w}'
        ws[ty] = f'=INDEX(账户_类型号,{a})'
        pn = norm(f'INDEX(账户_人员,{a})')
        ws[per] = f'=IF({ty}=6,{unit_match(pn, "简称", pn)},"")'
        first = f'{C8["首"]}{w}'
        ws[first] = (f'=IF(AND({ty}=6,{per}<>""),IF(COUNTIFS(${C8["类型"]}$2:{ty},6,${C8["人"]}$2:{per},{esc(per)})=1,1,0),0)')
        op = (f'SUMIFS(期初_借,期初_科目号,{k122103},期初_明细,{esc(per)},期初_有效,1)'
              f'-SUMIFS(期初_贷,期初_科目号,{k122103},期初_明细,{esc(per)},期初_有效,1)')
        lk = f'SUMIFS(资_净额,资_原对方号,{k122103},资_账户类型,"<>6",资_单位,{esc(per)})'
        ws[f'{C8["余额"]}{w}'] = (f'=IF({ty}<>6,0,ROUND(SUMIFS(资_净额,资_账户号,{a},资_有效,1)+IF({first}=1,{op}-{lk},0),2))')
    # 槽位：1801 付款 N8A、⑧1801 N_OP、1801 其余 1、1123 付款 N8A、⑧1123 N_OP、1123 其余 1、备用金 N_ACC、122103 其余 1、其他科目 N_COA
    Z8 = hx.cols(['s', '科目号', '明细', '额', '月数', '摊进', '说明', '出', '累'])
    for k, col in Z8.items():
        _lab(ws, f'{col}1', '⑧槽' + k)
    segs = [('付', N8A), ('期初', N_OP), ('1801余', 1), ('1123付', N8A), ('1123期初', N_OP), ('1123余', 1),
            ('备用金', N_ACC), ('122103余', 1), ('科目', N_COA)]
    s = 0
    sa = lambda col: f'${col}$2:${col}${N8A + 1}'
    sb = lambda col: f'${col}$2:${col}${N_OP + 1}'
    sc = lambda col: f'${C8[col]}$2:${C8[col]}${N_ACC + 1}'
    tgt_txt = lambda t_: f'IF({t_}={K("560207")},"管理",IF({t_}={K("410106")},"制造",""))'
    for seg, cnt in segs:
        for i in range(1, cnt + 1):
            s += 1
            w = s + 1
            ws[f'{Z8["s"]}{w}'] = s
            Z = lambda k: f'{Z8[k]}{w}'
            if seg == '付':
                rw = i + 1
                ws[Z('科目号')] = f'={K("1801")}'
                ws[Z('明细')] = f'=IF({A8["行"]}{rw}=0,"",INDEX(资_显示摘要,{A8["行"]}{rw})&"")'
                ws[Z('额')] = f'={A8["剩额"]}{rw}'
                ws[Z('月数')] = f'={A8["剩月"]}{rw}'
                ws[Z('摊进')] = '=' + tgt_txt(f'{A8["目标"]}{rw}')
                ws[Z('说明')] = f'=IF({A8["行"]}{rw}=0,"","上年结转：资金台帐第 "&{A8["行"]}{rw}&" 条的长期待摊")'
            elif seg == '期初':
                rw = i + 1
                ws[Z('科目号')] = f'={K("1801")}'
                ws[Z('明细')] = f'=INDEX(期初_明细,{i})&""'
                ws[Z('额')] = f'={B8["剩额"]}{rw}'
                ws[Z('月数')] = f'={B8["剩月"]}{rw}'
                ws[Z('摊进')] = '=' + tgt_txt(f'{B8["目标"]}{rw}')
                ws[Z('说明')] = '="上年结转：⑧ 期初的长期待摊"'
            elif seg == '1801余':
                ws[Z('科目号')] = f'={K("1801")}'
                ws[Z('明细')] = '=""'
                ws[Z('额')] = f'=ROUND(INDEX(矩_汇年末,{K("1801")})-SUM({sa(A8["剩额"])})-SUM({sb(B8["剩额"])}),2)'   # 含下级（1801 家族不进「科目」段）
                ws[Z('月数')] = '=0'
                ws[Z('摊进')] = '=""'
                ws[Z('说明')] = '="⚠ 手工凭证等记的长期待摊（没有摊销计划），剩余月数、摊进请自己填"'
            elif seg == '1123付':
                rw = i + 1
                t_ = f'{A8b["目标"]}{rw}'
                ws[Z('科目号')] = f'={K("1123")}'
                ws[Z('明细')] = f'=IF({A8b["行"]}{rw}=0,"",INDEX(资_显示摘要,{A8b["行"]}{rw})&"")'
                ws[Z('额')] = f'={A8b["剩额"]}{rw}'
                ws[Z('月数')] = f'={A8b["剩月"]}{rw}'
                ws[Z('摊进')] = f'=IF({t_}=0,"",INDEX(科目_编码,{t_})&" "&INDEX(科目_名称,{t_}))'
                ws[Z('说明')] = f'=IF({A8b["行"]}{rw}=0,"","上年结转：资金台帐第 "&{A8b["行"]}{rw}&" 条的预付（新本接着按月摊）")'
            elif seg == '1123期初':
                rw = i + 1
                ws[Z('科目号')] = f'=INDEX(期初_科目号,{i})'
                ws[Z('明细')] = f'=INDEX(期初_明细,{i})&""'
                ws[Z('额')] = f'={B8b["剩额"]}{rw}'
                ws[Z('月数')] = f'={B8b["剩月"]}{rw}'
                ws[Z('摊进')] = f'={B8b["摊进"]}{rw}'
                ws[Z('说明')] = (f'=IF({B8b["可"]}{rw}=1,"上年结转：⑧ 期初的预付（新本接着按月摊）",'
                                f'"⚠ 上年 ⑧ 的预付没摊（摊进、剩余月数没填或认不出来）：要摊就补上，不摊就是普通预付")')
            elif seg == '1123余':
                k1123 = K('1123')
                ws[Z('科目号')] = f'={k1123}'
                ws[Z('明细')] = '=""'
                ws[Z('额')] = (f'=IF({k1123}=0,0,ROUND(INDEX(矩_汇年末,{k1123})-SUM(${A8b["剩额"]}$2:${A8b["剩额"]}${N8A + 1})'
                              f'-SUM(${B8b["剩额"]}$2:${B8b["剩额"]}${N_OP + 1}),2))')
                ws[Z('月数')] = '=0'
                ws[Z('摊进')] = '=""'
                ws[Z('说明')] = '="⚠ 手工凭证等记的预付（没有摊销计划）：要摊就填剩余月数、摊进，不摊就是普通预付"'
            elif seg == '备用金':
                rw = i + 1
                ws[Z('科目号')] = f'={k122103}'
                ws[Z('明细')] = f'={C8["人"]}{rw}&""'
                ws[Z('额')] = f'={C8["余额"]}{rw}'
                ws[Z('月数')] = '=0'
                ws[Z('摊进')] = '=""'
                ws[Z('说明')] = '="上年结转：备用金"'
            elif seg == '122103余':
                ws[Z('科目号')] = f'={k122103}'
                ws[Z('明细')] = '=""'
                ws[Z('额')] = f'=ROUND(INDEX(矩_年末,{k122103})-SUM({sc("余额")}),2)'
                ws[Z('月数')] = '=0'
                ws[Z('摊进')] = '=""'
                ws[Z('说明')] = '="上年结转：备用金（对不上人的）"'
            else:
                no = Z('科目号')
                ws[no] = f'=IFERROR(MOD(SMALL(矩_排序键,{i}),1000),0)'
                k1 = f'INDEX(科目_一级号,MAX(1,{no}))'
                bad = 'OR(' + ','.join([f'{k1}={K(c)}' for c in FAM8] + [f'{no}={K(c)}' for c in EXACT8]) + ')'
                ws[Z('额')] = (f'=IF({no}=0,0,IF(OR(INDEX(科目_类别号,{no})=5,{bad}),0,INDEX(矩_年末,{no})))')
                ws[Z('明细')] = (f'=IF({no}=0,"",IF(AND(COUNTIF(期初_科目号,{no})=1,COUNTIFS(期初_科目号,{no},期初_有效,1)=1),'
                                f'INDEX(期初_明细,MATCH({no},期初_科目号,0))&"",""))')
                ws[Z('月数')] = '=0'
                ws[Z('摊进')] = '=""'
                ws[Z('说明')] = '="上年结转"'
            ws[Z('出')] = f'=IF(ROUND({Z("额")},2)<>0,1,0)'
            ws[Z('累')] = f'={Z("出")}' + (f'+{Z8["累"]}{w - 1}' if s > 1 else '')
    N8S = s
    zr = lambda k: f'${Z8[k]}$2:${Z8[k]}${N8S + 1}'
    r1 = sec(r, '③ 【基础资料】⑧ 其他期初余额：长期待摊、预付（每笔一行，带剩余月数、摊进）、制造费用留存、备用金（按人）、其他科目余额', OP_COLS, 'G')
    for j in range(1, N8_OUT + 1):
        rr = r1 + j - 1
        ws[f'{YN}{rr}'] = j
        ws[f'X{rr}'] = f'=IFERROR(MATCH(${YN}{rr},{zr("累")},0),0)'
        ws[f'X{rr}'].font = F_HELP
        S_ = f'$X{rr}'
        g = f'{S_}=0'
        no = f'INDEX({zr("科目号")},{S_})'
        amt = f'INDEX({zr("额")},{S_})'
        ws[f'A{rr}'] = f'=IF({g},"",INDEX(科目_编码,{no})&" "&INDEX(科目_名称,{no}))'
        ws[f'B{rr}'] = f'=IF({g},"",INDEX({zr("明细")},{S_}))'
        ws[f'C{rr}'] = f'=IF({g},"",IF({amt}>0,{amt},""))'
        ws[f'D{rr}'] = f'=IF({g},"",IF({amt}<0,-{amt},""))'
        ws[f'E{rr}'] = f'=IF({g},"",IF(INDEX({zr("月数")},{S_})>0,INDEX({zr("月数")},{S_}),""))'
        ws[f'F{rr}'] = f'=IF({g},"",INDEX({zr("摊进")},{S_}))'
        ws[f'G{rr}'] = f'=IF({g},"",INDEX({zr("说明")},{S_}))'
    style_rows(ws, r1, r1 + N8_OUT - 1, list('ABCDEFG'), auto=tuple('ABCDEFG'), fmts={'C': MONEY, 'D': MONEY, 'E': '0'},
               aligns={'A': AL, 'B': AL, 'G': AL})
    rr = r1 + N8_OUT
    ws.merge_cells(f'A{rr}:G{rr}')
    tot8 = f'SUM(C{r1}:C{rr - 1})-SUM(D{r1}:D{rr - 1})'
    n8 = f'MAX({zr("累")})'
    put(ws, f'A{rr}', (f'="⑧ 一共 "&{n8}&" 行"&IF({n8}>{N8_OUT},"　✗ 超过 {N8_OUT} 行，后面的没列出来，找我加",'
                       f'IF({n8}>{N_OP},"　⚠ 新本 ⑧ 只有 {N_OP} 行，先在新本 ⑧ 里插够行再粘",""))'
                       f'&"　　合计（借－贷）："&FIXED({tot8},2)&"（不带损益类、利润分配，新本未分配利润倒挤）"'),
        F_NOTE, fill('FFF2F2F2'), align=ALW, border=False)
    ws.row_dimensions[rr].height = 30
    chk_cells.append(f'A{rr}')
    r = rr + 2

    # ── ④ 往来单位 ──
    r1 = sec(r, '④ 【往来单位】整表：期末应收、应付、没开票、欠票（出借人的不填，在 ⑦）', UNIT_COLS, YE_LAST)
    for i in range(1, N_UNIT + 1):
        rr = r1 + i - 1
        ws[f'{YN}{rr}'] = i
        n_ = f'${YN}{rr}'
        g = f'INDEX(位_简称,{n_})=""'
        lend = f'INDEX(位_出借人类别,{n_})>0'
        for k, col in UNIT_COLS.items():
            if k in ('期初应收', '期初应付', '期初未开票', '期初欠票'):
                fld = {'期初应收': '应收', '期初应付': '应付', '期初未开票': '未开票', '期初欠票': '欠票'}[k]
                v = f'ROUND(INDEX(位_{fld},{n_}),2)'
                stop = f'OR({g},{lend})' if fld in ('应收', '应付') else g
                ws[f'{col}{rr}'] = f'=IF({stop},"",IF({v}=0,"",{v}))'
            else:
                ws[f'{col}{rr}'] = ye_copy(inref(SH_UNIT, col), n_, g)
    ucols = list(UNIT_COLS.values())
    style_rows(ws, r1, r1 + N_UNIT - 1, ucols, auto=tuple(ucols),
               fmts={UNIT_COLS['期初应收']: MONEY, UNIT_COLS['期初应付']: MONEY, UNIT_COLS['期初未开票']: MONEY,
                     UNIT_COLS['期初欠票']: MONEY, UNIT_COLS['到期日']: DATE},
               aligns={UNIT_COLS['简称']: AL, UNIT_COLS['全称']: AL, UNIT_COLS['备注']: AL})
    rr = r1 + N_UNIT
    H_, I_ = UNIT_COLS['期初应收'], UNIT_COLS['期初应付']
    r = check_row(rr, YE_LAST, [('应收（应收账款 1122）', f'SUM({H_}{r1}:{H_}{rr - 1})', bal('1122')),
                                ('应付（应付账款 2202）', f'SUM({I_}{r1}:{I_}{rr - 1})', f'-{bal("2202")}')])

    # ── ⑤ 物料档案 原料 ──
    r1 = sec(r, '⑤ 【物料档案】① 原料：12 月底的数量、单价、金额（进货即耗用的没有库存）', RM_COLS, 'O')
    q12 = lambda nm, k: f'INDEX({nm},{k},12)'
    for k in range(1, N_RM + 1):
        rr = r1 + k - 1
        ws[f'{YN}{rr}'] = k
        k = f'${YN}{rr}'
        g = f'INDEX(料档_名称,{k})=""'
        for f_, col in RM_COLS.items():
            if f_ == '期初数量':
                ws[f'{col}{rr}'] = f'=IF(OR({g},ROUND({q12("料_末量", k)},4)=0),"",{q12("料_末量", k)})'
            elif f_ == '期初金额':
                ws[f'{col}{rr}'] = (f'=IF(OR({g},AND(ROUND({q12("料_末量", k)},4)=0,ROUND({q12("料_末额", k)},2)=0)),"",'
                                    f'{q12("料_末额", k)})')
            elif f_ == '期初单价':
                ws[f'{col}{rr}'] = (f'=IF(OR({g},ROUND({q12("料_末量", k)},4)=0),"",'
                                    f'ROUND({q12("料_末额", k)}/{q12("料_末量", k)},6))')
            else:
                ws[f'{col}{rr}'] = ye_copy(inref(SH_ITEM, col, block='原料'), k, g)
    rc = list(RM_COLS.values())
    style_rows(ws, r1, r1 + N_RM - 1, rc, auto=tuple(rc),
               fmts={RM_COLS['期初金额']: MONEY, RM_COLS['期初单价']: '#,##0.00####', RM_COLS['期初数量']: '#,##0.000',
                     RM_COLS['含量']: '0.0%'}, aligns={RM_COLS['名称']: AL, RM_COLS['备注']: AL})
    rr = r1 + N_RM
    Lc = RM_COLS['期初金额']
    r = check_row(rr, 'O', [('原材料 1403', f'SUM({Lc}{r1}:{Lc}{rr - 1})', bal('1403'))])

    # ── ⑥ 物料档案 产品 ──
    r1 = sec(r, '⑥ 【物料档案】② 产品：12 月底的数量、单价、金额，在产（生产成本-直接材料）', PRD_COLS, 'M')
    for p in range(1, N_PRD + 1):
        rr = r1 + p - 1
        ws[f'{YN}{rr}'] = p
        p = f'${YN}{rr}'
        g = f'INDEX(品档_名称,{p})=""'
        for f_, col in PRD_COLS.items():
            if f_ == '期初数量':
                ws[f'{col}{rr}'] = f'=IF(OR({g},ROUND({q12("品_末量", p)},4)=0),"",{q12("品_末量", p)})'
            elif f_ == '期初金额':
                ws[f'{col}{rr}'] = (f'=IF(OR({g},AND(ROUND({q12("品_末量", p)},4)=0,ROUND({q12("品_末额", p)},2)=0)),"",'
                                    f'{q12("品_末额", p)})')
            elif f_ == '期初单价':
                ws[f'{col}{rr}'] = (f'=IF(OR({g},ROUND({q12("品_末量", p)},4)=0),"",'
                                    f'ROUND({q12("品_末额", p)}/{q12("品_末量", p)},6))')
            elif f_ == '期初在产':
                ws[f'{col}{rr}'] = f'=IF(OR({g},ROUND({q12("本_末在产", p)},2)=0),"",{q12("本_末在产", p)})'
            else:
                ws[f'{col}{rr}'] = ye_copy(inref(SH_ITEM, col, block='产品'), p, g)
    pc = list(PRD_COLS.values())
    style_rows(ws, r1, r1 + N_PRD - 1, pc, auto=tuple(pc),
               fmts={PRD_COLS['期初金额']: MONEY, PRD_COLS['期初单价']: '#,##0.00####', PRD_COLS['期初数量']: '#,##0.000',
                     PRD_COLS['期初在产']: MONEY}, aligns={PRD_COLS['名称']: AL, PRD_COLS['备注']: AL})
    rr = r1 + N_PRD
    Kc, Lp = PRD_COLS['期初金额'], PRD_COLS['期初在产']
    r = check_row(rr, 'M', [('库存商品 1405', f'SUM({Kc}{r1}:{Kc}{rr - 1})', bal('1405')),
                            ('在产 400101', f'SUM({Lp}{r1}:{Lp}{rr - 1})', bal('400101'))])

    # ── ⑦ 固定资产 ──
    r1 = sec(r, '⑦ 【固定资产】卡片：建账前已提＝到 12 月底的累计折旧（今年处置的不带）', FA_COLS, 'K')
    for i in range(1, N_FA + 1):
        rr = r1 + i - 1
        ws[f'{YN}{rr}'] = i
        i = f'${YN}{rr}'
        g = f'OR(INDEX(卡_名称,{i})="",INDEX(卡_处置月,{i})<=12)'
        for f_, col in FA_COLS.items():
            if f_ == '建账前已提':
                ws[f'{col}{rr}'] = (f'=IF({g},"",IF(INDEX(卡_有效,{i})=1,ROUND(INDEX(卡_建账前已提,{i})+SUM(INDEX(卡_月折旧块,{i},0)),2),'
                                    f'IF(INDEX({inref(SH_FA, col)},{i}+1)="","",INDEX({inref(SH_FA, col)},{i}+1))))')
            else:
                ws[f'{col}{rr}'] = ye_copy(inref(SH_FA, col), i, g)
    fc = list(FA_COLS.values())
    style_rows(ws, r1, r1 + N_FA - 1, fc, auto=tuple(fc),
               fmts={FA_COLS['原值']: MONEY, FA_COLS['建账前已提']: MONEY, FA_COLS['取得日期']: DATE,
                     FA_COLS['处置日期']: DATE, FA_COLS['残值率']: '0.0%'}, aligns={FA_COLS['名称']: AL, FA_COLS['备注']: AL})
    rr = r1 + N_FA
    Fc, Ic = FA_COLS['原值'], FA_COLS['建账前已提']
    r = check_row(rr, 'K', [('固定资产 1601', f'SUM({Fc}{r1}:{Fc}{rr - 1})', bal('1601')),
                            ('累计折旧 1602', f'SUM({Ic}{r1}:{Ic}{rr - 1})', f'-{bal("1602")}')])
    # 辅助区、X 列（槽位号）藏起来
    for i in range(CI(HX0), hx.c + 1):
        ws.column_dimensions[CL(i)].hidden = True
    ws.column_dimensions['X'].hidden = True
    ws.column_dimensions[YN].hidden = True
    for rr in range(1, ws.max_row + 1):
        if ws[f'{YN}{rr}'].value is not None:
            ws[f'{YN}{rr}'].font = F_HELP
    home_link(ws, 'T1')
    ws.column_dimensions['T'].width = 10
    ws.freeze_panes = 'B4'
    print_setup(ws, '1:2', landscape=True)
    ws.print_area = f'A1:{YE_LAST}{r}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    return chk_cells


def build(wb, ctx=None):
    build_fen(wb[H_FEN])
    build_pz(wb[H_PZ])
    build_ju(wb, wb[H_JU])
    build_vlist(wb, wb[SH_VLIST])
    build_vprint(wb, wb[SH_VPRINT])
    build_tb(wb, wb[SH_TB])
    build_ye(wb, wb[SH_YE])
    for sh in (SH_VLIST, SH_VPRINT, SH_TB, SH_YE):
        protect(wb[sh])
