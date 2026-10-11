# -*- coding: utf-8 -*-
"""A072 化工 · 引擎隐藏表（月份 1～12，建账月以前全 0）：
   _料（原料进销存，月末一次加权）、_品（产品进销存）、_本（生产成本：在产、直接材料、制造费用池按明细各自分配）、
   _息（借款/投资按日单利、每月计提）、_摊（1801/1123 按月摊）、_工汇（工资按部门×月）。

   规矩（防 Excel 静态循环）：
   - 逐月递推只引用同一行前一个月那一格（相邻列），不经 INDEX/OFFSET；
   - 一个字段一块 12 列（layout.eng_block）；按月份生成公式时 SUMIF/INDEX 只框某个字段块的「那一列」；
   - 只对隐藏层（_资 _进 _销 _产 _工 _票 _手 _盘）、基础镜像和别的引擎块求和，不对 _分 求和。
   依赖顺序（没有回路）：_本 槽投/废料扣 ← 镜像；_料 ← 槽投、废料扣；_本 分额/DM ← _料 投额；
   制造费用池 ← _料 盘额、_卡、_摊、_工汇、镜像；_本 OH/完工 ← 池；_品 ← _本 完工。

   月末自动计提只做 m≤P_截止月：_摊、_息 本月利息、_工汇、盘点（_料/_品 盘量）、制造费用 本月分配。
   _料/_品/_本 的数量、投料、完工、销售不挡截止月（截止月以后的月份没分制造费用，池留着）。

   私有块/列（别的模块不要引用，名称只用 layout.NAMES 和本模块加的 息_*、摊_贷1801、摊_贷1123）：
   _料：HX 起 可用量、可用额、出量、出额、运比（公开字段后空一列，各 12 列）。
   _品：FD 起 可用量、可用额（同上）。
   _本：DG i、DH 制费明细号（池行标签）；FQ i、FR p（分到行标签）；
        槽投块（第 (j−1)×N_PRD＋p＋1 行：j 槽、p 产品；GF j、GG p、GH 槽料号，GI..GT 槽投、GU..HF 分额）。
        槽和原料一一对应（_物 槽_料号：一个槽只认唯一一种料，一种料只占一个槽），所以「料投」就是槽投那一行；
        _料 投量＝SUMIF(槽料号列, k, 槽投第 m 列)，分额按槽内各产品累计舍入，合计正好等于 _料 投额。
   _息：BX..CG 标量（名称 息_<标量>，第 2..N_LEND+1 行）；CI 单位号、CJ 借日率、CK 投日率；⑦ 辅助 CM n、CN 金额×日期（第 n 条在第 n+1 行）。
   _摊：第 N_FEE+3 行 贷1801、N_FEE+4 行 贷1123（C..N 12 列，名称 摊_贷1801、摊_贷1123）；
        AB..AI ⑧ 期初 1801 辅助（第 n 条在第 n+1 行：n、额、月数、目标号、始月、止月、月摊、尾差）。"""
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import F_HELP

MS = range(1, 13)


class Eng:
    """字段块列号：公开字段从 col0 起每个 12 列；私有字段在公开字段后空 gap 列起"""

    def __init__(self, fields, col0='C', priv=(), gap=1):
        self.pos = {}
        c = CI(col0)
        for f in fields:
            self.pos[f] = c
            c += 12
        c += gap
        for f in priv:
            self.pos[f] = c
            c += 12
        self.end = c - 1

    def col(self, f, m):
        return CL(self.pos[f] + m - 1)

    def c(self, f, m, r):
        return f'{self.col(f, m)}{r}'

    def rng(self, f, m, r1, r2, sh=None):
        cl = self.col(f, m)
        return (f'{sh}!' if sh else '') + f'${cl}${r1}:${cl}${r2}'


def _lab(ws, a, t):
    ws[a] = t
    ws[a].font = F_HELP


def _heads(ws, E, fields):
    for f in fields:
        for m in MS:
            _lab(ws, f'{E.col(f, m)}1', f'{f}{m}')


def rec(m, init, prev):
    """递推：建账月以前 0，建账月＝初值，之后＝前一个月那一格"""
    if m == 1:
        return f'IF(1<P_建账月,0,{init})'
    return f'IF({m}<P_建账月,0,IF({m}=P_建账月,{init},{prev}))'


def _name(wb, nm, ref):
    wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


# ═══════════════════ 列布局（模块内共享） ═══════════════════
LIAO_PRIV = ['可用量', '可用额', '出量', '出额', '运比']
PIN_PRIV = ['可用量', '可用额']
EL = Eng(LIAO_FIELDS, priv=LIAO_PRIV)
EP = Eng(PIN_FIELDS, priv=PIN_PRIV)
EB = Eng(BEN_PRD_FIELDS)                                  # 产品块
EO = Eng(BEN_OH_FIELDS, col0=BEN_OH_COL0)                 # 制造费用池块
OH_I, OH_NO = CL(CI(BEN_OH_COL0) - 2), CL(CI(BEN_OH_COL0) - 1)        # DG i、DH 制费明细号
SH_I, SH_P = CL(CI(BEN_SH_COL0) - 2), CL(CI(BEN_SH_COL0) - 1)         # FQ i、FR p
SH_C0 = CI(BEN_SH_COL0)                                   # 分到 第 m 月＝SH_C0+m−1
_s0 = SH_C0 + 12 + 1
SL_J, SL_P, SL_K = CL(_s0), CL(_s0 + 1), CL(_s0 + 2)      # 槽投块标签：j、p、槽料号
ES = Eng(['槽投', '分额'], col0=CL(_s0 + 3))
N_SL = N_SLOT * N_PRD
EX = Eng(XI_FIELDS)
XI_SC0 = EX.end + 1                                       # 标量从 BX 起（Eng.end 已含空的一列 BW）
XI_SC = {k: CL(XI_SC0 + i) for i, k in enumerate(XI_SCALARS)}
_x0 = XI_SC0 + len(XI_SCALARS) + 1
XI_PRIV = dict(单位号=CL(_x0), 借日率=CL(_x0 + 1), 投日率=CL(_x0 + 2))
XI_LO = dict(n=CL(_x0 + 4), 金额日=CL(_x0 + 5))           # ⑦ 辅助：第 n 条在第 n+1 行
ET = Eng(TAN_FIELDS)
TAN_ROW = {'贷1801': N_FEE + 3, '贷1123': N_FEE + 4}
_t0 = ET.end + 1
TAN_OP = {k: CL(_t0 + i) for i, k in enumerate(['n', '额', '月数', '目标号', '始月', '止月', '月摊', '尾差'])}
EG = Eng(GH_FIELDS)

R_RM, R_PRD, R_FEE = N_RM + 1, N_PRD + 1, N_FEE + 1


# ═══════════════════ _本 槽投（先排：_料 要用） ═══════════════════
def build_ben(ws):
    B = H_BEN
    _lab(ws, 'A1', 'p')
    _lab(ws, 'B1', '产品')
    _heads(ws, EB, BEN_PRD_FIELDS)
    _lab(ws, f'{OH_I}1', 'i')
    _lab(ws, f'{OH_NO}1', '制费明细号')
    _heads(ws, EO, BEN_OH_FIELDS)
    _lab(ws, f'{SH_I}1', 'i')
    _lab(ws, f'{SH_P}1', 'p')
    for m in MS:
        _lab(ws, f'{CL(SH_C0 + m - 1)}1', f'分到{m}')
    for a, t in ((SL_J, '槽j'), (SL_P, '产品p'), (SL_K, '槽料号')):
        _lab(ws, f'{a}1', t)
    _heads(ws, ES, ['槽投', '分额'])
    W = lambda m, r1, r2: f'${EB.col("权重", m)}${r1}:${EB.col("权重", m)}${r2}'
    # ── 产品块 ──
    for p in range(1, N_PRD + 1):
        r = p + 1
        ws[f'A{r}'] = p
        ws[f'B{r}'] = f'=INDEX(品档_名称,$A{r})'
        c = lambda f, m: EB.c(f, m, r)
        rk = f'INDEX(品档_回收料号,$A{r})'
        for m in MS:
            key = f'$A{r}*100+{m}'
            ws[c('产量', m)] = f'=SUMIF(产_品月键,{key},产_成品吨)'
            ws[c('权重', m)] = f'={c("产量", m)}*INDEX(品档_系数,$A{r})'
            ws[c('初在产', m)] = '=' + rec(m, f'INDEX(品档_期初在产,$A{r})', c('末在产', m - 1) if m > 1 else '')
            ws[c('DM', m)] = (f'=SUMIF(${SL_P}$2:${SL_P}${N_SL + 1},$A{r},{ES.rng("分额", m, 2, N_SL + 1)})')
            ws[c('废料扣', m)] = (f'=ROUND(SUMIF(产_品月键,{key},产_废料吨)*IF({rk}>0,INDEX(料档_定额,MAX(1,{rk})),0),2)')
            ws[c('转出DM', m)] = (f'=IF({c("产量", m)}>0,ROUND({c("初在产", m)}+{c("DM", m)}-{c("废料扣", m)},2),0)')
            fs = CL(SH_C0 + m - 1)
            ws[c('OH', m)] = f'=SUMIF(${SH_P}$2:${SH_P}${N_OH * N_PRD + 1},$A{r},${fs}$2:${fs}${N_OH * N_PRD + 1})'
            ws[c('完工', m)] = f'=ROUND({c("转出DM", m)}+{c("OH", m)},2)'
            ws[c('末在产', m)] = (f'=ROUND({c("初在产", m)}+{c("DM", m)}-{c("废料扣", m)}-{c("转出DM", m)},2)')
    # ── 制造费用池（第 i 个 4101 明细）──
    ins = 'IFERROR(INDEX(类别_付号,MATCH("员工保险",类别_名称,0)),0)'
    for i in range(1, N_OH + 1):
        r = i + 1
        no = f'${OH_NO}{r}'
        ws[f'{OH_I}{r}'] = i
        ws[no[1:]] = f'=INDEX(制费明细_号,${OH_I}{r})'
        c = lambda f, m: EO.c(f, m, r)
        op = (f'IF({no}=0,0,ROUND(SUMIFS(期初_借,期初_科目号,{no},期初_有效,1)'
              f'-SUMIFS(期初_贷,期初_科目号,{no},期初_有效,1),2))')
        for m in MS:
            ws[c('上月留', m)] = '=' + rec(m, op, c('留下月', m - 1) if m > 1 else '')
            pan = EL.rng('盘额', m, 2, R_RM, H_LIAO)
            ka = f'{H_KA}!${CL(CI(KA_MON) + m - 1)}$2:${CL(CI(KA_MON) + m - 1)}${N_FA + 1}'
            t18 = ET.rng('摊1801', m, 2, R_FEE, H_TAN)
            t23 = ET.rng('摊1123', m, 2, R_FEE, H_TAN)
            ga = lambda f, d: f'{H_GH}!{EG.c(f, m, d + 1)}'
            sb = f'SUM({H_GH}!{EG.c("扣社保", m, 2)}:{EG.c("扣社保", m, 5)})'
            parts = [
                f'SUMIFS(资_费用额,资_原对方号,{no},资_费用月,{m},资_分摊,0)',
                f'+SUMIFS(进_入账额,进_借方号,{no},进_月,{m})',
                f'+SUMIFS(手_借,手_科目号,{no},手_月,{m},手_记账,1)-SUMIFS(手_贷,手_科目号,{no},手_月,{m},手_记账,1)',
                f'-IF({no}={kh("410199")},SUM({pan}),0)',
                f'+IF({no}={kh("410107")},SUMIF(卡_部门号,1,{ka}),0)',
                f'+SUMIF(费用明细_号,{no},{t18})+SUMIF(费用明细_号,{no},{t23})',
                f'+IF({no}={kh("410104")},{ga("应发", 1)},0)+IF({no}={kh("410105")},{ga("应发", 2)},0)',
                f'-IF({no}={ins},{sb},0)',
                f'-SUMIFS(票_补票额,票_补票目标号,{no},票_月,{m})',
            ]
            ws[c('本月发生', m)] = f'=IF(OR({no}=0,{m}<P_建账月),0,ROUND({"".join(parts)},2))'
            ws[c('可分配', m)] = f'=ROUND({c("上月留", m)}+{c("本月发生", m)},2)'
            ws[c('本月分配', m)] = f'=IF(AND({m}<=P_截止月,SUM({W(m, 2, R_PRD)})>0),{c("可分配", m)},0)'
            ws[c('留下月', m)] = f'=ROUND({c("可分配", m)}-{c("本月分配", m)},2)'
    # ── 分到（第 (i−1)×N_PRD＋p＋1 行）：累计舍入 ──
    for i in range(1, N_OH + 1):
        for p in range(1, N_PRD + 1):
            r = (i - 1) * N_PRD + p + 1
            ws[f'{SH_I}{r}'] = i
            ws[f'{SH_P}{r}'] = p
            for m in MS:
                A = EO.c('本月分配', m, i + 1)
                tot = f'SUM({W(m, 2, R_PRD)})'
                t1 = f'ROUND({A}*SUM({W(m, 2, p + 1)})/{tot},2)'
                t0 = f'-ROUND({A}*SUM({W(m, 2, p)})/{tot},2)' if p > 1 else ''
                ws[f'{CL(SH_C0 + m - 1)}{r}'] = f'=IF({A}=0,0,{t1}{t0})'
    # ── 槽投、分额（第 (j−1)×N_PRD＋p＋1 行）──
    for j in range(1, N_SLOT + 1):
        g0, g1 = (j - 1) * N_PRD + 2, j * N_PRD + 1
        for p in range(1, N_PRD + 1):
            r = (j - 1) * N_PRD + p + 1
            ws[f'{SL_J}{r}'] = j
            ws[f'{SL_P}{r}'] = p
            ws[f'{SL_K}{r}'] = f'=N(INDEX(槽_料号,${SL_J}{r}))'
            for m in MS:
                q = ES.col('槽投', m)
                ws[f'{q}{r}'] = f'=SUMIF(产_品月键,${SL_P}{r}*100+{m},产_槽{j})'
                T = f'INDEX({EL.rng("投额", m, 2, R_RM, H_LIAO)},MAX(1,${SL_K}{r}))'
                S = f'SUM(${q}${g0}:${q}${g1})'
                t1 = f'ROUND({T}*SUM(${q}${g0}:${q}{r})/{S},2)'
                t0 = f'-ROUND({T}*SUM(${q}${g0}:${q}{r - 1})/{S},2)' if p > 1 else ''
                ws[ES.c('分额', m, r)] = f'=IF(OR(${SL_K}{r}=0,{S}=0),0,{t1}{t0})'


# ═══════════════════ _料 ═══════════════════
def build_liao(ws):
    _lab(ws, 'A1', 'k')
    _lab(ws, 'B1', '原料')
    _heads(ws, EL, LIAO_FIELDS + LIAO_PRIV)
    sk = f'{H_BEN}!${SL_K}$2:${SL_K}${N_SL + 1}'
    for k in range(1, N_RM + 1):
        r = k + 1
        ws[f'A{r}'] = k
        ws[f'B{r}'] = f'=INDEX(料档_名称,$A{r})'
        c = lambda f, m: EL.c(f, m, r)
        A = f'$A{r}'
        p0 = f'IF(INDEX(料档_期初单价,{A})>0,INDEX(料档_期初单价,{A}),INDEX(料档_参考单价,{A}))'
        for m in MS:
            key = f'{A}*100+{m}'
            pv = lambda f: c(f, m - 1) if m > 1 else ''
            ws[c('期初量', m)] = '=' + rec(m, f'INDEX(料档_期初量,{A})', pv('末量'))
            ws[c('期初额', m)] = '=' + rec(m, f'INDEX(料档_期初额,{A})', pv('末额'))
            ws[c('期初运', m)] = '=' + rec(m, '0', pv('末运'))
            ws[c('进量', m)] = f'=SUMIF(进_料月键,{key},进_数量)'
            ws[c('进额', m)] = f'=SUMIF(进_料月键,{key},进_入账额)'
            ws[c('进运', m)] = f'=SUMIF(进_料月键,{key},进_运费入账)'
            ws[c('回量', m)] = f'=SUMIF(产_回月键,{key},产_废料吨)'
            fk = EB.col('废料扣', m)
            ws[c('回额', m)] = f'=SUMIF(品档_回收料号,{A},{H_BEN}!${fk}$2:${fk}${R_PRD})'
            ws[c('投量', m)] = f'=SUMIF({sk},{A},{ES.rng("槽投", m, 2, N_SL + 1, H_BEN)})'
            ws[c('销量', m)] = f'=SUMIFS(销_吨位,销_出库键,{key},销_出库,2)'
            ws[c('可用量', m)] = f'={c("期初量", m)}+{c("进量", m)}+{c("回量", m)}'
            ws[c('可用额', m)] = f'=ROUND({c("期初额", m)}+{c("进额", m)}+{c("回额", m)},2)'
            ws[c('出量', m)] = f'={c("投量", m)}+{c("销量", m)}'
            q, a, o = c('可用量', m), c('可用额', m), c('出量', m)
            back = p0 if m == 1 else f'IF({m}=P_建账月,{p0},{c("单价", m - 1)})'
            ws[c('单价', m)] = f'=IF({m}<P_建账月,0,IF(ROUND({q},4)>0,ROUND({a}/{q},6),{back}))'
            full = f'AND(ROUND({q},4)>0,ROUND({o}-{q},4)=0)'
            u = c('单价', m)
            ws[c('出额', m)] = f'=IF({full},{a},ROUND({o}*{u},2))'
            ws[c('销额', m)] = f'=IF({full},ROUND({c("出额", m)}*{c("销量", m)}/{o},2),ROUND({c("销量", m)}*{u},2))'
            ws[c('投额', m)] = f'=ROUND({c("出额", m)}-{c("销额", m)},2)'
            ws[c('运比', m)] = f'=IF({a}<>0,({c("期初运", m)}+{c("进运", m)})/{a},0)'
            ws[c('投运', m)] = f'=ROUND({c("投额", m)}*{c("运比", m)},2)'
            ws[c('盘量', m)] = (f'=IF(AND({m}<=P_截止月,COUNTIFS(盘_键,{key},盘_重复,0)>0),'
                              f'ROUND(SUMIFS(盘_实盘,盘_键,{key},盘_重复,0)-({q}-{o}),6),0)')
            ws[c('盘额', m)] = f'=ROUND({c("盘量", m)}*{u},2)'
            ws[c('末量', m)] = f'={q}-{o}+{c("盘量", m)}'
            ws[c('末额', m)] = f'=ROUND({a}-{c("出额", m)}+{c("盘额", m)},2)'
            ws[c('末运', m)] = (f'=ROUND({c("期初运", m)}+{c("进运", m)}-ROUND({c("出额", m)}*{c("运比", m)},2)'
                              f'+ROUND({c("盘额", m)}*{c("运比", m)},2),2)')


# ═══════════════════ _品 ═══════════════════
def build_pin(ws):
    _lab(ws, 'A1', 'p')
    _lab(ws, 'B1', '产品')
    _heads(ws, EP, PIN_FIELDS + PIN_PRIV)
    for p in range(1, N_PRD + 1):
        r = p + 1
        ws[f'A{r}'] = p
        ws[f'B{r}'] = f'=INDEX(品档_名称,$A{r})'
        c = lambda f, m: EP.c(f, m, r)
        A = f'$A{r}'
        for m in MS:
            key = f'{A}*100+{m}'
            pv = lambda f: c(f, m - 1) if m > 1 else ''
            ws[c('期初量', m)] = '=' + rec(m, f'INDEX(品档_期初量,{A})', pv('末量'))
            ws[c('期初额', m)] = '=' + rec(m, f'INDEX(品档_期初额,{A})', pv('末额'))
            ws[c('完量', m)] = f'=SUMIF(产_品月键,{key},产_成品吨)'
            ws[c('完额', m)] = f'={H_BEN}!{EB.c("完工", m, r)}'
            ws[c('购量', m)] = f'=SUMIF(进_品月键,{key},进_数量)'
            ws[c('购额', m)] = f'=SUMIF(进_品月键,{key},进_入账额)'
            ws[c('销量', m)] = f'=SUMIFS(销_吨位,销_出库键,{key},销_出库,1)'
            ws[c('可用量', m)] = f'={c("期初量", m)}+{c("完量", m)}+{c("购量", m)}'
            ws[c('可用额', m)] = f'=ROUND({c("期初额", m)}+{c("完额", m)}+{c("购额", m)},2)'
            q, a, o = c('可用量', m), c('可用额', m), c('销量', m)
            p0 = f'INDEX(品档_期初单价,{A})'
            back = p0 if m == 1 else f'IF({m}=P_建账月,{p0},{c("单价", m - 1)})'
            ws[c('单价', m)] = f'=IF({m}<P_建账月,0,IF(ROUND({q},4)>0,ROUND({a}/{q},6),{back}))'
            u = c('单价', m)
            ws[c('销额', m)] = f'=IF(AND(ROUND({q},4)>0,ROUND({o}-{q},4)=0),{a},ROUND({o}*{u},2))'
            ws[c('盘量', m)] = (f'=IF(AND({m}<=P_截止月,COUNTIFS(盘_品键,{key},盘_重复,0)>0),'
                              f'ROUND(SUMIFS(盘_实盘,盘_品键,{key},盘_重复,0)-({q}-{o}),6),0)')
            ws[c('盘额', m)] = f'=ROUND({c("盘量", m)}*{u},2)'
            ws[c('末量', m)] = f'={q}-{o}+{c("盘量", m)}'
            ws[c('末额', m)] = f'=ROUND({a}-{c("销额", m)}+{c("盘额", m)},2)'


# ═══════════════════ _息 ═══════════════════
def build_xi(ws):
    _lab(ws, 'A1', 'l')
    _lab(ws, 'B1', '出借人')
    _heads(ws, EX, XI_FIELDS)
    for k, col in {**XI_SC, **XI_PRIV}.items():
        _lab(ws, f'{col}1', k)
    _lab(ws, f'{XI_LO["n"]}1', '⑦n')
    _lab(ws, f'{XI_LO["金额日"]}1', '⑦金额×日期')
    for n in range(1, N_LO + 1):
        r = n + 1
        ws[f'{XI_LO["n"]}{r}'] = n
        ws[f'{XI_LO["金额日"]}{r}'] = f'=INDEX(借期_金额,${XI_LO["n"]}{r})*INDEX(借期_日期,${XI_LO["n"]}{r})'
    lo_ad = f'${XI_LO["金额日"]}$2:${XI_LO["金额日"]}${N_LO + 1}'
    pd_rng = f'{H_ZI}!${ZI["本金乘日"]}$2:${ZI["本金乘日"]}${N_CASH + 1}'   # 直接写区域
    for l in range(1, N_LEND + 1):
        r = l + 1
        A = f'$A{r}'
        ws[f'A{r}'] = l
        ws[f'B{r}'] = f'=INDEX(出借_名称,{A})'
        u = f'${XI_PRIV["单位号"]}{r}'
        ws[u.replace('$', '')] = f'=N(INDEX(出借_单位号,{A}))'
        days = f'INDEX(位_计息天数,{u})'
        ws[f'{XI_PRIV["借日率"]}{r}'] = f'=IF({u}=0,0,INDEX(位_借款利率,{u})/{days})'
        ws[f'{XI_PRIV["投日率"]}{r}'] = f'=IF({u}=0,0,INDEX(位_投资利率,{u})/{days})'
        sc = lambda k: f'{XI_SC[k]}{r}'
        for pre, kind, crit in (('借', 1, '"<>2"'), ('投', 2, '2')):
            d = f'${XI_PRIV[pre + "日率"]}{r}'
            lo = f'借期_出借人号,{A},借期_性质,{crit},借期_有效,1'
            c = lambda f, m: EX.c(pre + f, m, r)
            for m in MS:
                key = f'{A}*1000+{kind * 100 + m}'
                init = f'SUMIFS(借期_金额,{lo})'
                prev = f'{c("月初本金", m - 1)}+{c("本月变动", m - 1)}' if m > 1 else ''
                ws[c('月初本金', m)] = '=' + rec(m, init, prev)
                ws[c('本月变动', m)] = f'=SUMIF(资_息键,{key},资_本金变动)'
                nx = f'DATE(P_年度,{m + 1},1)'
                st = f'MAX(DATE(P_年度,{m},1),P_建账日)'
                ws[c('本月利息', m)] = (f'=IF(OR({m}>P_截止月,{m}<P_建账月,{d}=0),0,ROUND({d}*({c("月初本金", m)}*({nx}-{st})'
                                      f'+{nx}*{c("本月变动", m)}-SUMIF(资_息键,{key},{pd_rng})),2))')
            ws[sc(pre + '累计利息')] = f'=SUM({c("本月利息", 1)}:{c("本月利息", 12)})'
            ws[sc(pre + '期初应付息')] = f'=ROUND(SUMIFS(借期_应计息,{lo})-SUMIFS(借期_已付息,{lo}),2)'
            code = '2231' if kind == 1 else '2232'
            ws[sc(pre + '已付息')] = (f'=IF({u}=0,0,-ROUND(SUMIFS(资_净额,资_对方号,{kh(code)},资_单位号,{u},'
                                     f'资_日期,"<="&P_截止),2))')
            zi = f'资_出借人号,{A},资_性质,{kind},资_日期,"<="&P_截止'
            ws[sc(pre + '本金')] = f'=ROUND(SUMIFS(借期_金额,{lo})+SUMIFS(资_本金变动,{zi}),2)'
            pd = f'SUMIFS({lo_ad},{lo})+SUMIFS({pd_rng},{zi})'
            ws[sc(pre + '到截止应计')] = (f'=ROUND({d}*((P_截止+1)*{sc(pre + "本金")}-({pd})),2)'
                                       f'-SUMIFS(借期_已付息,{lo})-{sc(pre + "已付息")}')


# ═══════════════════ _摊 ═══════════════════
def build_tan(ws):
    _lab(ws, 'A1', 'j')
    _lab(ws, 'B1', '费用明细号')
    _heads(ws, ET, TAN_FIELDS)
    for k, col in TAN_OP.items():
        _lab(ws, f'{col}1', '⑧' + k)
    O = TAN_OP
    R = lambda k: f'${O[k]}$2:${O[k]}${N_OP + 1}'
    for n in range(1, N_OP + 1):
        r = n + 1
        x = lambda f: f'INDEX(期初_{f},${O["n"]}{r})'
        ws[f'{O["n"]}{r}'] = n
        ok = (f'AND({x("有效")}=1,{x("科目号")}>0,INDEX(科目_一级号,MAX(1,{x("科目号")}))={kh("1801")})')
        ws[f'{O["额"]}{r}'] = f'=IF({ok},ROUND({x("借")}-{x("贷")},2),0)'
        ws[f'{O["月数"]}{r}'] = f'=IF({x("剩余月数")}>0,{x("剩余月数")},P_摊销月数)'
        ws[f'{O["目标号"]}{r}'] = f'=IF({O["额"]}{r}=0,0,IF({x("摊进")}=2,{kh("560207")},{kh("410106")}))'
        ws[f'{O["始月"]}{r}'] = '=MAX(P_建账月,P_投产月)'
        ws[f'{O["止月"]}{r}'] = f'={O["始月"]}{r}+{O["月数"]}{r}-1'
        ws[f'{O["月摊"]}{r}'] = f'=IF({O["额"]}{r}=0,0,ROUND({O["额"]}{r}/{O["月数"]}{r},2))'
        ws[f'{O["尾差"]}{r}'] = f'=IF({O["额"]}{r}=0,0,ROUND({O["额"]}{r}-{O["月摊"]}{r}*{O["月数"]}{r},2))'
    for j in range(1, N_FEE + 1):
        r = j + 1
        ws[f'A{r}'] = j
        ws[f'B{r}'] = f'=INDEX(费用明细_号,$A{r})'
        no = f'$B{r}'
        for m in MS:
            off = f'OR({no}=0,{m}<P_建账月,{m}>P_截止月)'
            for f, kind in (('摊1801', 1), ('摊1123', 2)):
                z = (f'SUMIFS(资_月摊额,资_分摊,{kind},资_摊目标号,{no},资_摊始月,"<={m}",资_摊止月,">={m}")'
                     f'+SUMIFS(资_摊尾差,资_分摊,{kind},资_摊目标号,{no},资_摊止月,{m})')
                if kind == 1:
                    z += (f'+SUMIFS({R("月摊")},{R("目标号")},{no},{R("始月")},"<={m}",{R("止月")},">={m}")'
                          f'+SUMIFS({R("尾差")},{R("目标号")},{no},{R("止月")},{m})')
                ws[ET.c(f, m, r)] = f'=IF({off},0,ROUND({z},2))'
    for f, src in (('贷1801', '摊1801'), ('贷1123', '摊1123')):
        r = TAN_ROW[f]
        _lab(ws, f'A{r}', f)
        for m in MS:
            ws[f'{CL(CI("C") + m - 1)}{r}'] = f'=ROUND(SUM({ET.rng(src, m, 2, R_FEE)}),2)'


# ═══════════════════ _工汇 ═══════════════════
def build_gh(ws):
    _lab(ws, 'A1', 'd')
    _lab(ws, 'B1', '部门')
    _heads(ws, EG, GH_FIELDS)
    src = dict(应发='工_应发', 扣社保='工_扣社保', 其他扣款='工_其他扣款', 实发='工_实发')
    for d in range(1, 5):
        r = d + 1
        ws[f'A{r}'] = d
        ws[f'B{r}'] = DEPTS[d - 1]
        for f in GH_FIELDS:
            for m in MS:
                ws[EG.c(f, m, r)] = (f'=IF(OR({m}<P_建账月,{m}>P_截止月),0,ROUND(SUMIF(工_部门月键,$A{r}*100+{m},{src[f]}),2))')


def define_names(wb):
    for k, col in XI_SC.items():
        _name(wb, f'息_{k}', f'{H_XI}!${col}$2:${col}${N_LEND + 1}')
    for f, r in TAN_ROW.items():
        _name(wb, f'摊_{f}', f'{H_TAN}!$C${r}:$N${r}')


def build(wb, ctx=None):
    define_names(wb)
    build_ben(wb[H_BEN])
    build_liao(wb[H_LIAO])
    build_pin(wb[H_PIN])
    build_xi(wb[H_XI])
    build_tan(wb[H_TAN])
    build_gh(wb[H_GH])
