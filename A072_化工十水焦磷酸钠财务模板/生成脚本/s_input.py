# -*- coding: utf-8 -*-
"""A072 化工 · 录入表（蓝）：资金台帐、进货物流登记、销售台账、生产日报、工资表、发票登记、月末盘点、手工凭证（会计组）；
   基础（灰）：基础资料（① 参数 ②～⑧）、往来单位、物料档案（① 原料 ② 产品）、固定资产、固定支出、会计科目。

   录入表只放录入列，最右边几列是只给人看的显示列：=IFERROR(INDEX(隐藏列,ROW()-ROW($A$4)),"") 一类（别的公式不引用它们，
   插进来的行没有这几格也不影响汇总，第 3 行有「右边没公式」提示）。列一律按 layout 的 *_COLS / *_SHOW。
   预置数据全部来自 ctx（data.build_ctx()）；写法跟 t/s_raw.py 一样（'YYYY-MM-DD' 变日期），空的不写。
   基础资料 ②～⑧ 的块标题必须跟 layout.BASE_BLOCKS 一字不差写在 A 列（s_mirror 用 MATCH 找块的边界）；块里不放别的字，
   容量提示写在右边说明里。"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font
from common import *
from layout import *          # layout 后导入：norm/esc/颜色以 layout 为准

FILL_EX = fill('FFEDEDED')     # 示例行
F_LBL = fill('FFD9E1F2')
H_AUTO = 'FF8EA9DB'            # 自动列表头
H_SLOT = 'FF808080'            # 生产日报原料表头（公式，别改）
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC65911')
F_INFO = Font(name=YH, sz=10, color='FF2F75B5')
MON_FMT = '[<=12]0"月";[>=36526]yyyy-mm;0'
QTY = '#,##0.###;[Red]-#,##0.###;"-"'
PRICE = '#,##0.00##'
EX = '示例'


# ───────────────────────── 小工具 ─────────────────────────
def _v(x):
    """跟 s_raw 一样：'YYYY-MM-DD' → 日期；空串、None → 不写"""
    if x is None or (isinstance(x, str) and x.strip() == ''):
        return None
    if isinstance(x, (dt.date, dt.datetime)):
        return dt.datetime(x.year, x.month, x.day)
    if isinstance(x, str) and len(x) == 10 and x[4] == '-' and x[7] == '-':
        try:
            return dt.datetime.strptime(x, '%Y-%m-%d')
        except ValueError:
            return x
    return x


def _t(s):
    assert s is None or len(s) <= 255, f'有效性文字超过 255 字（{len(s)}）：{s[:40]}…'
    return s


def _add(ws, dv, sqref, prompt=None):
    if prompt:
        dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', _t(prompt), True
    _t(dv.error)
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def dv_names(ws, sqref, name, prompt=None, stop=True, error=None):
    """下拉：来源是 layout.DV_NAMES 的名称（formula1 不带 =）"""
    assert name in DV_NAMES, name
    dv = DataValidation(type='list', formula1=name, allow_blank=True, showErrorMessage=True,
                        errorStyle='stop' if stop else 'warning', errorTitle='不在清单里',
                        error=error or ('请从下拉里选。客户、供应商、物流、出借人先到【往来单位】加一行；账户、收支类别、报销项目、人员、'
                                        '票种先到【基础资料】加一行；物料、产品先到【物料档案】加一行'))
    return _add(ws, dv, sqref, prompt)


def dv_fixed(ws, sqref, items, prompt=None, stop=True, error=None):
    dv = DataValidation(type='list', formula1='"' + ','.join(items) + '"', allow_blank=True, showErrorMessage=True,
                        errorStyle='stop' if stop else 'warning', errorTitle='请从下拉里选',
                        error=error or ('只能填：' + '、'.join(items)))
    return _add(ws, dv, sqref, prompt)


def dv_dec(ws, sqref, what='金额', lo='-999999999', hi='999999999', extra='', prompt=None):
    dv = DataValidation(type='decimal', operator='between', formula1=lo, formula2=hi, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what,
                        error=f'{what}要填数字（不要带文字、单位）{extra}')
    return _add(ws, dv, sqref, prompt)


def dv_whole(ws, sqref, lo, hi, what, prompt=None):
    dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what, error=f'{what}填 {lo}～{hi} 的整数')
    return _add(ws, dv, sqref, prompt)


def dv_tip(ws, sqref, prompt):
    """只给提示、不限制输入"""
    dv = DataValidation(allow_blank=True, showErrorMessage=False)
    return _add(ws, dv, sqref, prompt)


def dv_day(ws, sqref, prompt=None):
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle='日期',
                        error='输 2026-8-1 或 2026/8/1 就行（不要用 8.1，那是小数）')
    return _add(ws, dv, sqref, prompt)


def dv_len(ws, sqref, n=30):
    dv = DataValidation(type='textLength', operator='lessThanOrEqual', formula1=str(n), allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle='太长', error=f'不要超过 {n} 个字')
    return _add(ws, dv, sqref)


def _hd(k, req=(), sub=''):
    """表头：必填带 ★；第二行小字说明"""
    return (f'{k}★' if k in req else k) + (f'\n（{sub}）' if sub else '')


def _grey(ws, r, c1, c2):
    for c in range(CI(c1), CI(c2) + 1):
        ws.cell(row=r, column=c).fill = FILL_EX


def _is_ex(note):
    return str(note or '').startswith(EX)


def _frame(ws, last_col, heads, color, tip, W, nrows, fm, al, freeze, auto=(), title_col=None, hdr_h=44, filt=True):
    widths(ws, W)
    title(ws, ws.title, title_col or last_col, color, tip)
    header(ws, HDR, heads, color, height=hdr_h)
    style_rows(ws, R0, R0 + nrows - 1, [c for c, _ in heads], auto=auto, fmts=fm, aligns=al)
    ws.freeze_panes = freeze
    if filt:
        ws.auto_filter.ref = f'A{HDR}:{last_col}{R0 + nrows - 1}'
    for col in auto:
        ws[f'{col}{HDR}'].fill = fill(H_AUTO)


def _write_dicts(ws, rows, cols, r0=R0):
    """录入表：ctx 每行是 dict（键＝layout 列名）；生产日报的「槽」是 12 个原料吨数的列表"""
    for i, x in enumerate(rows):
        r = r0 + i
        for k, v in x.items():
            if k == '槽':
                for j, q in enumerate(v):
                    q = _v(q)
                    if q is not None:
                        ws[f'{PROD_SLOT_COLS[j]}{r}'] = q
                continue
            if k not in cols:
                continue
            v = _v(v)
            if v is not None:
                ws[f'{cols[k]}{r}'] = v
        if _is_ex(x.get('备注')):
            last_in = max(cols.values(), key=CI)
            if ws.title == SH_PROD:
                last_in = PROD_COLS['备注']
            _grey(ws, r, 'A', last_in)


def _write_tuples(ws, rows, cols, r0, note_key=None):
    """基础表：ctx 每行是元组，顺序＝layout 列顺序"""
    keys = list(cols)
    for i, row in enumerate(rows):
        r = r0 + i
        for k, v in zip(keys, row):
            v = _v(v)
            if v is not None:
                ws[f'{cols[k]}{r}'] = v
        if note_key and _is_ex(dict(zip(keys, row)).get(note_key)):
            _grey(ws, r, 'A', max(cols.values(), key=CI))


def _print_to_last(ws, last_col, par_key, cap):
    """打印区域只到最后一行有内容的（_参 里算的行号），不打几千行空表格线"""
    q = f"'{ws.title}'"
    cell = PAR[par_key]
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${last_col}$1:${last_col}${R0 + cap - 1},'
                                f'MIN({R0 + cap - 1},{H_PAR}!${cell[0]}${cell[1:]}))')
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'


def _kpi(ws, row, cells):
    """第 3 行：标签（合并两格）＋ 值"""
    for c, lbl, f, fmt in cells:
        put(ws, f'{c}{row}', lbl, F_KPI_L, F_LBL, align=AC)
        ws.merge_cells(f'{c}{row}:{CL(CI(c) + 1)}{row}')
        put(ws, f'{CL(CI(c) + 2)}{row}', f, F_AUTOB, FILL_AUTO, fmt, AR)


def _label(ws, rng, text):
    c1 = rng.split(':')[0]
    put(ws, c1, text, F_KPI_L, F_LBL, align=AC)
    if ':' in rng:
        ws.merge_cells(rng)


def _sub(ws, col, last, fmt=MONEY):
    put(ws, f'{col}3', f'=SUBTOTAL(9,{col}{HDR}:{col}{last})', F_AUTOB, FILL_AUTO, fmt, AR)


def _show(ws, nrows, cols_formulas, fmts=None, aligns=None, r0=R0, hdr=HDR):
    """显示列：{n}＝第 n 条（ROW()-ROW(表头)），{r}＝本行行号；整条包 IFERROR，出错显示空"""
    n = f'ROW()-ROW($A${hdr})'
    fmts, aligns = fmts or {}, aligns or {}
    for i in range(nrows):
        r = r0 + i
        for col, f in cols_formulas.items():
            c = ws[f'{col}{r}']
            c.value = '=IFERROR(' + f.replace('{n}', n).replace('{r}', str(r)) + ',"")'
            c.fill = FILL_AUTO
            c.font = F_AUTO
            if col in fmts:
                c.number_format = fmts[col]
            c.alignment = aligns.get(col, AC)


def _check_cf(ws, col, r0, last):
    rng = f'{col}{r0}:{col}{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{r0},1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{r0},1)="⚠"'], font=F_WARN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{r0},1)="ℹ"'], font=F_INFO))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{r0}="✓"'], font=Font(name=YH, sz=10, color='FF70AD47')))


def _nofx(ws, key_col, chk_col, cell, last, merge_to=None):
    """「右边没公式的行」提示：有内容、问题列却是空的（插进来的行没拉公式，或者超了容量）"""
    cnt = f'COUNTIFS({key_col}{R0}:{key_col}{last},"<>",{chk_col}{R0}:{chk_col}{last},"")'
    f = f'=IF({cnt}>0,{cnt}&" 行右边没公式（插进来的行，或者超了容量）：选上一行右边灰色几格往下拉","")'
    put(ws, cell, f, F_RED, align=AL, border=False)
    if merge_to:
        ws.merge_cells(f'{cell}:{merge_to}')


def _chk(pre):
    """这一行的问题：✗ 一条链，⚠/ℹ 一条链；有 ✗ 先显示 ✗"""
    e, w = f'INDEX({pre}_错,{{n}})&""', f'INDEX({pre}_提醒,{{n}})&""'
    return f'IF(INDEX({pre}_末行,{{n}})=0,"",IF({e}<>"",{e},IF({w}<>"",{w},"✓")))'


def _chk1(pre):
    """只有一列「校验」的隐藏表（_位 _固 _卡）"""
    return f'IF(INDEX({pre}_末行,{{n}})=0,"",IF(INDEX({pre}_校验,{{n}})&""="","✓",INDEX({pre}_校验,{{n}})&""))'


def own(col, hdr=HDR):
    """显示列里取本行的一格：按第 n 条取（插行、删行、排序都不乱）"""
    return f'INDEX(${col}${hdr}:${col}${REF_END},{{n}}+1)'


def own_row(c1, c2, hdr=HDR):
    """本行 c1:c2 一段（COUNTA 用）"""
    return f'INDEX(${c1}${hdr}:${c2}${REF_END},{{n}}+1,0)'


def _tail(ws, last_col, par_key, cap):
    home_link(ws, f'{CL(CI(last_col) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, last_col, par_key, cap)


def rng_of(cap, r0=R0):
    return lambda col: f'{col}{r0}:{col}{r0 + cap - 1}'


# ───────────────────────── 资金台帐 ─────────────────────────
def build_cash(ws, ctx):
    C, S = CASH_COLS, CASH_SHOW
    sub = {'往来单位': '收付货款、借款、还垫付必选', '编号': '垫付必填；固定支出填编号', '报销人': '员工报销必填',
           '报销项目': '员工报销必填', '票种': '空＝默认', '所属月份': '电费等填用的月份；空＝付款月'}
    heads = [(C[k], _hd(k, CASH_REQ, sub.get(k, ''))) for k in C] + [
        (S['余额'], '余额\n（本账户，按上下顺序）'), (S['单位类型'], '单位类型\n（自动）'), (S['算到哪'], '算到哪\n（自动）'),
        (S['记到'], '记到\n（自动，会计科目）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 所有账户的每一笔钱都记这里，一笔一行，收入、支出只填一个。收客户货款、付货款（原材料材料款、原材料运输费、销货运费）、'
           '借款、还款一定要选往来单位。老板自己掏钱付的：账户选「老板垫付」，类别选买了什么，编号填第 3 行的「下一个垫付编号」，'
           '票种空着＝待补票（补到票在【发票登记】登）；公司还老板的钱只在银行记一行，类别选「垫付现金已报」、往来单位选老板，'
           '垫付账户那边不用再记。员工报销：类别选员工报销，报销人、报销项目必填。发备用金在银行记一行（类别发备用金、往来单位选拿钱的人），'
           '员工花备用金在他的备用金账户记支出。缴税要选是哪种税（增值税／附加税印花税／所得税／代扣个税），'
           '一张缴款书几种税拆几行。电费、水费、供热这种先用后付的，所属月份填用的那个月（只对费用生效）。内部转账记两行（转出、转进）。'
           '右边灰色几列自动出：账户余额、单位类型、算到哪、记到哪个科目、这一行的问题；可以按它们筛选，第 3 行是筛出来的合计。灰底是示例行，正式用之前删掉。')
    W = {C['日期']: 11, C['账户']: 11, C['收支类别']: 14, C['往来单位']: 13, C['摘要']: 22, C['收入']: 12, C['支出']: 12,
         C['编号']: 12, C['报销人']: 9, C['报销项目']: 10, C['票种']: 9, C['所属月份']: 11, C['备注']: 18,
         S['余额']: 13, S['单位类型']: 9, S['算到哪']: 10, S['记到']: 24, S['校验']: 40}
    fm = {C['日期']: DATE, C['收入']: MONEY, C['支出']: MONEY, C['编号']: '@', C['所属月份']: MON_FMT}
    al = {C[k]: AL for k in ('往来单位', '摘要', '备注', '编号')}
    al.update({C['收入']: AR, C['支出']: AR})
    rows = ctx.get('cash', [])
    _frame(ws, CASH_LAST, heads, C_IN, tip, W, N_CASH, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=C['备注'])
    last = R0 + N_CASH + 2000
    _label(ws, 'A3:E3', '筛选后看得见的行合计 →')
    _sub(ws, C['收入'], last)
    _sub(ws, C['支出'], last)
    _label(ws, f'{C["编号"]}3', '下一个垫付编号')
    put(ws, f'{C["报销人"]}3', '=P_下一垫付号', F_KPI_V, FILL_AUTO, '@', AC)
    ws.merge_cells(f'{C["报销人"]}3:{C["报销项目"]}3')
    _label(ws, f'{C["票种"]}3', '收入－支出')
    put(ws, f'{C["所属月份"]}3', f'={C["收入"]}3-{C["支出"]}3', F_AUTOB, FILL_AUTO, MONEY, AR)
    _nofx(ws, 'A', S['校验'], f'{C["备注"]}3', last, f'{S["校验"]}3')
    _write_dicts(ws, rows, C)
    end = 'INDEX(资_末行,{n})=0'
    uno = 'INDEX(资_单位号,{n})'
    unit = 'INDEX(资_单位,{n})&""'
    person = f'ISNUMBER(MATCH({esc(unit)},人员_姓名,0))'
    cat = 'INDEX(资_类别号,{n})'
    k1 = 'INDEX(资_原对方一级,{n})'
    kind = f'IF({cat}>0,INDEX(类别_归类,{cat})&"","")'
    re_kind = (f'IF({k1}={kh("4101")},"生产成本",IF({k1}={kh("5601")},"销售费用",IF({k1}={kh("5602")},"管理费用",'
               f'IF({k1}={kh("5603")},"财务费用",{kind}))))')
    acct = 'INDEX(资_原对方号,{n})'
    _show(ws, N_CASH, {
        S['余额']: f'IF(OR({end},INDEX(资_余额,{{n}})&""=""),"",INDEX(资_余额,{{n}}))',
        S['单位类型']: (f'IF(OR({end},{unit}=""),"",IF({uno}>0,INDEX(位_类型,{uno})&"",IF(OR(INDEX(资_报销人号,{{n}})>0,{person}),"员工",'
                     f'"没建档")))'),
        S['算到哪']: f'IF({end},"",IF(AND(INDEX(资_报销项目号,{{n}})>0,{k1}>0),{re_kind},{kind}))',
        S['记到']: (f'IF(OR({end},{acct}<=0),"",INDEX(科目_全称,{acct})'
                  f'&IF(INDEX(资_跨月,{{n}})=1,"（归"&INDEX(资_所属月,{{n}})&"月）",IF(INDEX(资_跨月,{{n}})=2,"（建账前的）",""))'
                  f'&IF(INDEX(资_分摊,{{n}})>0,"（分月摊）","")&IF(N(INDEX(资_税额,{{n}}))<>0,"＋进项税",""))'),
        S['校验']: _chk('资'),
    }, {S['余额']: MONEY}, {S['余额']: AR, S['记到']: AL, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_CASH - 1)
    rng = rng_of(N_CASH)
    dv_day(ws, rng(C['日期']))
    dv_names(ws, rng(C['账户']), '账户列表', '从【基础资料】②账户里选。老板个人掏钱替公司付的，账户选「老板垫付」')
    dv_names(ws, rng(C['收支类别']), '收支类别列表',
             '从【基础资料】③里选。付进货物流登记里欠的货款、运费选「原材料材料款」「原材料运输费」「销货运费」；还老板垫付选「垫付现金已报」；'
             '缴税选是哪种税')
    dv_names(ws, rng(C['往来单位']), '往来单位列表',
             '收付货款、借款还款、还老板垫付要选【往来单位】里的；零星的（菜市场、加油站）可以直接打字', stop=False)
    dv_tip(ws, rng(C['编号']), '老板垫付的每一笔都填编号：年月-X序号（202608-X01），第 3 行有下一个编号；付【固定支出】里的款填那条的编号（GD01）')
    dv_names(ws, rng(C['报销人']), '人员列表', '员工报销、备用金：从【基础资料】⑤人员里选')
    dv_names(ws, rng(C['报销项目']), '报销项目列表', '员工报销：从【基础资料】④里选（材料配件、加油、日常用品…），科目、票种按它算')
    dv_names(ws, rng(C['票种']), '票种列表',
             '空着：老板垫付、备用金账户＝待补票；员工报销＝报销项目的默认；其他＝收支类别的默认。费用、设备的专票才抵进项')
    dv_tip(ws, rng(C['所属月份']), '电费、水费、供热、工资这种先用后付的，填用的那个月（1～12）。只对费用类生效：付货款、还借款等填了不算。'
                                   '空着＝付款那个月')
    for k in ('收入', '支出'):
        dv_dec(ws, rng(C[k]), extra='；10000+200 要写成 =10000+200')
    _tail(ws, CASH_LAST, '资金末行', N_CASH)


# ───────────────────────── 进货物流登记 ─────────────────────────
def build_pur(ws, ctx):
    C, S = PUR_COLS, PUR_SHOW
    sub = {'日期': '到货', '物料': '原料、辅料、吨袋…', '数量': '按物料的单位', '单价': '空＝待定价',
           '货款金额': '空＝数量×单价', '货款票种': '空＝供应商默认', '物流单位': '运费付给谁；空＝供应商',
           '运费': '空＝数量×每吨运费', '税点': '空＝运费×7%，无票不加', '运费金额': '空＝运费＋税点', '运费票种': '空＝物流默认'}
    heads = [(C[k], _hd(k, PUR_REQ, sub.get(k, ''))) for k in C] + [
        (S['货款'], '货款\n（自动；待定价按参考价暂估）'), (S['运费合计'], '运费合计\n（自动，含税点）'),
        (S['到厂单价'], '到厂单价\n（自动，货款＋运费÷数量）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 进一车货登一行（照物流登记表），货款和运费在同一行。物料从下拉选：原料、辅料、吨袋都在【物料档案】里，买回来直接卖的成品也能选。'
           '数量按物料的单位（原料吨、吨袋个、活性炭 kg）。还没定价的，单价、货款金额空着＝按物料档案的参考单价暂估，定了价回原行补上。'
           '运费：填运费（按吨算的填每吨运费），税点空着自动按 7% 加（运费票种是无票的不加），运费金额＝运费＋税点；'
           '物流单位空着＝供应商送货、运费也付给供应商。扣款、扣杂：数量空着、货款金额填负数；退货：数量填负数。'
           '吨袋、活性炭这类也登这里，物流几列空着就行。付钱在【资金台帐】选「原材料材料款」「原材料运输费」。灰底是示例行。')
    W = {C['日期']: 11, C['车牌号']: 11, C['司机']: 8, C['联系电话']: 12, C['车型']: 12, C['装货点']: 11, C['卸货点']: 9,
         C['物料']: 17, C['供应商']: 13, C['数量']: 9, C['单价']: 9, C['货款金额']: 12, C['货款票种']: 9, C['物流单位']: 12,
         C['每吨运费']: 8, C['运费']: 10, C['税点']: 9, C['运费金额']: 11, C['运费票种']: 9, C['备注']: 20,
         S['货款']: 13, S['运费合计']: 12, S['到厂单价']: 12, S['校验']: 40}
    fm = {C['日期']: DATE, C['数量']: QTY, C['单价']: PRICE, C['货款金额']: MONEY, C['每吨运费']: PRICE, C['运费']: MONEY,
          C['税点']: MONEY, C['运费金额']: MONEY, C['联系电话']: '@', C['车牌号']: '@'}
    al = {C[k]: AL for k in ('车牌号', '车型', '装货点', '卸货点', '物料', '供应商', '物流单位', '备注', '联系电话')}
    for k in ('数量', '单价', '货款金额', '每吨运费', '运费', '税点', '运费金额'):
        al[C[k]] = AR
    rows = ctx.get('purchases', [])
    _frame(ws, PUR_LAST, heads, C_IN, tip, W, N_PUR, fm, al, f'B{R0}', auto=tuple(S.values()), title_col=C['备注'])
    last = R0 + N_PUR + 2000
    _label(ws, 'A3:I3', '筛选后看得见的行合计 →')
    _sub(ws, C['数量'], last, QTY)
    _label(ws, f'{C["单价"]}3', '待定价')
    put(ws, f'{C["货款金额"]}3', '=COUNTIFS(进_待定价,1,进_有效,1)&" 车（全部）"', F_AUTOB, FILL_AUTO, align=AR)
    _nofx(ws, 'A', S['校验'], f'{C["货款票种"]}3', last, f'{C["备注"]}3')
    _sub(ws, S['货款'], last)
    _sub(ws, S['运费合计'], last)
    _write_dicts(ws, rows, C)
    end = 'INDEX(进_末行,{n})=0'
    q = 'INDEX(进_数量,{n})'
    _show(ws, N_PUR, {
        S['货款']: f'IF({end},"",INDEX(进_货款计,{{n}}))',
        S['运费合计']: f'IF({end},"",INDEX(进_运费计,{{n}}))',
        S['到厂单价']: f'IF(OR({end},N({q})=0),"",ROUND((N(INDEX(进_货款计,{{n}}))+N(INDEX(进_运费计,{{n}})))/{q},2))',
        S['校验']: _chk('进'),
    }, {S['货款']: MONEY, S['运费合计']: MONEY, S['到厂单价']: PRICE},
        {S['货款']: AR, S['运费合计']: AR, S['到厂单价']: AR, S['校验']: AL})
    lastshow = R0 + N_PUR - 1
    _check_cf(ws, S['校验'], R0, lastshow)
    j, k, l = C['数量'], C['单价'], C['货款金额']
    ws.conditional_formatting.add(f'{S["货款"]}{R0}:{S["货款"]}{lastshow}', FormulaRule(
        formula=[f'AND(ISNUMBER(${j}{R0}),${j}{R0}>0,${k}{R0}="",${l}{R0}="")'], font=F_WARN, fill=cf_bg('FFFFF2CC')))
    rng = rng_of(N_PUR)
    dv_day(ws, rng(C['日期']))
    dv_fixed(ws, rng(C['车型']), ['13米高栏国五', '9.6米高栏', '17.5米平板', '4.2米厢式'], stop=False,
             error='不在常用车型里（可以直接写，点「是」）')
    dv_names(ws, rng(C['物料']), '物料产品列表',
             '从【物料档案】里选（原料、辅料、吨袋…；买回来直接卖的成品也在这里）。新料先去物料档案加一行')
    dv_names(ws, rng(C['供应商']), '供应商列表', '从【往来单位】的供应商里选；新供应商先去往来单位加一行')
    dv_dec(ws, rng(C['数量']), '数量', prompt='按物料的单位：原料吨、吨袋个、活性炭 kg（过磅数可以有小数）。扣款、扣杂：数量空着，货款金额填负数；退货：数量填负数')
    dv_dec(ws, rng(C['单价']), '单价', prompt='含税单价。还没定价就空着（按物料档案的参考单价暂估），定了回来补上')
    dv_dec(ws, rng(C['货款金额']), '货款金额', prompt='空着＝数量×单价；跟磅单、结算单对不上就直接填金额')
    dv_names(ws, rng(C['货款票种']), '票种列表', '空着＝供应商在【往来单位】里的默认票种')
    dv_names(ws, rng(C['物流单位']), '往来单位列表',
             '运费付给谁（物流公司、个人司机都要在【往来单位】建档，类型选物流）。空着＝供应商送货、运费也付给供应商',
             error='物流单位要先在【往来单位】加一行（类型选物流）')
    dv_dec(ws, rng(C['每吨运费']), '每吨运费', prompt='按吨算运费的填每吨多少；按车一口价的空着，直接填运费')
    dv_dec(ws, rng(C['运费']), '运费', prompt='一车的运费（不含税点）；空着＝数量×每吨运费')
    dv_dec(ws, rng(C['税点']), '税点', prompt='空着＝运费×7%（基础资料里改）；运费票种是无票的不加。不加税点填 0')
    dv_dec(ws, rng(C['运费金额']), '运费金额', prompt='要付给物流的总数＝运费＋税点；空着自动算。只填了这一格也行（倒推运费）')
    dv_names(ws, rng(C['运费票种']), '票种列表', '空着＝物流单位的默认票种（个人司机在往来单位里设成无票）')
    _tail(ws, PUR_LAST, '进货末行', N_PUR)


# ───────────────────────── 销售台账 ─────────────────────────
def build_sale(ws, ctx):
    C, S = SALE_COLS, SALE_SHOW
    sub = {'日期': '合同日期', '客户': '销售单位', '产品': '空＝默认产品', '吨位': '吨', '金额': '这一车；空＝单价×吨位',
           '提货日期': '按它记收入；空＝日期', '签收对账': '填签收或日期', '发票号': '开了票就填', '票种': '空＝按客户',
           '物流单位': '我方出运费才填', '税点': '空＝运费×7%', '运费金额': '空＝运费＋税点', '运费票种': '空＝物流默认'}
    heads = [(C[k], _hd(k, SALE_REQ, sub.get(k, ''))) for k in C] + [
        (S['金额算'], '金额\n（自动）'), (S['结算情况'], '结算情况\n（自动，先发先收）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 发一车货登一行（一个合同几车就登几行，合同号、客户每行都填）。产品空着＝默认产品（十水焦钠）；卖废料、卖原料就选那种料。'
           '单价、金额填含税的（【基础资料】里可改）；还没定价的单价、金额空着，出库照算，定了价回原行补上。收入、出库按提货日期（空＝日期）。'
           '我方出运费的才填物流单位、运费，税点空着自动加 7%（无票不加）。开了票就在发票号里填号码，这一车就算已开票，不用再去发票登记登。'
           '客户扣款、折让：吨位空着、金额填负数；退货：吨位填负数。签收对账填「签收」或签收日期；结算情况按先发货先收钱自动算。灰底是示例行。')
    W = {C['日期']: 11, C['合同号']: 12, C['客户']: 13, C['产品']: 10, C['单价']: 9, C['吨位']: 8, C['金额']: 12,
         C['提货日期']: 11, C['车牌号']: 11, C['司机']: 8, C['司机电话']: 12, C['签收对账']: 10, C['发票号']: 21, C['票种']: 9,
         C['物流单位']: 11, C['运费']: 10, C['税点']: 9, C['运费金额']: 11, C['运费票种']: 9, C['备注']: 18,
         S['金额算']: 13, S['结算情况']: 18, S['校验']: 40}
    fm = {C['日期']: DATE, C['提货日期']: DATE, C['签收对账']: DATE, C['单价']: PRICE, C['吨位']: QTY, C['金额']: MONEY,
          C['运费']: MONEY, C['税点']: MONEY, C['运费金额']: MONEY, C['合同号']: '@', C['发票号']: '@', C['司机电话']: '@',
          C['车牌号']: '@'}
    al = {C[k]: AL for k in ('合同号', '客户', '车牌号', '发票号', '物流单位', '备注', '司机电话')}
    for k in ('单价', '吨位', '金额', '运费', '税点', '运费金额'):
        al[C[k]] = AR
    rows = ctx.get('sales', [])
    _frame(ws, SALE_LAST, heads, C_IN, tip, W, N_SALE, fm, al, f'D{R0}', auto=tuple(S.values()), title_col=C['备注'])
    last = R0 + N_SALE + 2000
    _label(ws, 'A3:E3', '筛选后看得见的行合计 →')
    _sub(ws, C['吨位'], last, QTY)
    _label(ws, f'{C["金额"]}3', '待定价')
    put(ws, f'{C["提货日期"]}3', '=COUNTIFS(销_待定价,1,销_有效,1)&" 车（全部）"', F_AUTOB, FILL_AUTO, align=AR)
    ws.merge_cells(f'{C["提货日期"]}3:{C["车牌号"]}3')
    _nofx(ws, C['客户'], S['校验'], f'{C["司机"]}3', last, f'{C["备注"]}3')
    _sub(ws, S['金额算'], last)
    _write_dicts(ws, rows, C)
    end = 'INDEX(销_末行,{n})=0'
    k = 'INDEX(销_客户号,{n})'
    amt = 'INDEX(销_收入含税,{n})'
    cum = f'SUMIFS(销_收入含税,销_客户号,{k},销_排序键,"<="&INDEX(销_排序键,{{n}}))'
    paid = f'(N(INDEX(位_回款,{k}))-N(INDEX(位_期初应收,{k})))'
    got = f'MAX(0,MIN({amt},{paid}-({cum}-{amt})))'
    sign = f'INDEX(${C["签收对账"]}${HDR}:${C["签收对账"]}${REF_END},{{n}}+1)'      # 本行的签收对账（按第 n 条取，插行排序不乱）
    due = f'(IF(ISNUMBER({sign}),{sign},WORKDAY(INDEX(销_日期,{{n}}),P_对账工作日))+P_付款天数)'
    settle = (f'IF(OR({end},N({k})<=0,N({amt})<=0,INDEX(销_日期,{{n}})>P_截止),"",'
              f'IF({got}>={amt}-0.005,"已结清",IF({got}>0.005,"部分（还差 "&TEXT({amt}-{got},"#,##0.00")&"）",'
              f'IF(P_截止>{due},"逾期 "&(P_截止-{due})&" 天","未到期（"&TEXT({due},"m/d")&"）"))))')
    _show(ws, N_SALE, {
        S['金额算']: f'IF({end},"",IF(INDEX(销_待定价,{{n}})=1,"待定价",INDEX(销_金额含税,{{n}})))',
        S['结算情况']: settle,
        S['校验']: _chk('销'),
    }, {S['金额算']: MONEY}, {S['金额算']: AR, S['结算情况']: AL, S['校验']: AL})
    lastshow = R0 + N_SALE - 1
    _check_cf(ws, S['校验'], R0, lastshow)
    a, st = S['金额算'], S['结算情况']
    ws.conditional_formatting.add(f'{a}{R0}:{a}{lastshow}', FormulaRule(formula=[f'{a}{R0}="待定价"'], font=F_WARN))
    ws.conditional_formatting.add(f'{st}{R0}:{st}{lastshow}', FormulaRule(formula=[f'LEFT({st}{R0},2)="逾期"'], font=F_RED))
    ws.conditional_formatting.add(f'{st}{R0}:{st}{lastshow}', FormulaRule(formula=[f'{st}{R0}="已结清"'],
                                                                         font=Font(name=YH, sz=10, color='FF70AD47')))
    rng = rng_of(N_SALE)
    for kk in ('日期', '提货日期'):
        dv_day(ws, rng(C[kk]))
    dv_names(ws, rng(C['客户']), '客户列表', '从【往来单位】的客户里选；新客户先去往来单位加一行')
    dv_names(ws, rng(C['产品']), '物料产品列表', '空着＝默认产品（十水焦钠）；卖废料、卖原料就选那种料')
    dv_dec(ws, rng(C['单价']), '单价', prompt='每吨含税单价（基础资料「销售单价含税」改成否就填不含税的）。还没定价就空着')
    dv_dec(ws, rng(C['吨位']), '吨位', prompt='这一车多少吨（过磅数可以有小数）。退货填负数；客户扣款、折让：吨位空着、金额填负数')
    dv_dec(ws, rng(C['金额']), '金额', prompt='这一车的金额；空着＝单价×吨位。还没定价的单价、金额都空着')
    dv_tip(ws, rng(C['签收对账']), '客户签收了填「签收」或签收日期。按签收日期（没填日期就按提货日期）往后几个工作日对账，再过几天到期（【基础资料】参数里改）')
    dv_tip(ws, rng(C['发票号']), '这一车开了票就填发票号码（数电票 20 位，这一列是文本格式直接输），就算已开票，不用再去【发票登记】登')
    dv_names(ws, rng(C['票种']), '票种列表', '空着＝按【往来单位】里这家客户的开票设置（要＝专票13%，不要＝无票）')
    dv_names(ws, rng(C['物流单位']), '往来单位列表', '我方出运费的才填：运费付给谁（要在【往来单位】建档）。客户自提不填',
             error='物流单位要先在【往来单位】加一行（类型选物流）')
    dv_dec(ws, rng(C['运费']), '运费', prompt='我方出的一车运费（不含税点）')
    dv_dec(ws, rng(C['税点']), '税点', prompt='空着＝运费×7%（基础资料里改）；运费票种是无票的不加。不加税点填 0')
    dv_dec(ws, rng(C['运费金额']), '运费金额', prompt='要付给物流的总数＝运费＋税点；空着自动算')
    dv_names(ws, rng(C['运费票种']), '票种列表', '空着＝物流单位的默认票种')
    _tail(ws, SALE_LAST, '销售末行', N_SALE)


# ───────────────────────── 生产日报 ─────────────────────────
def build_prod(ws, ctx):
    C, S = PROD_COLS, PROD_SHOW
    heads = [(C['日期'], _hd('日期', PROD_REQ)), (C['产品'], _hd('产品', PROD_REQ, '空＝默认产品')),
             (C['产量'], _hd('产量', PROD_REQ, '袋')), (C['废料'], _hd('废料', PROD_REQ, '袋，二次过滤')),
             (C['备注'], '备注'), (S['成品吨'], '成品吨\n（自动）'), (S['投料合计'], '投料合计\n（吨，自动）'),
             (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 每天每个产品登一行：产量、废料按袋（1 袋几吨在【物料档案】产品里填），投的原料按吨填在对应那一列。'
           '第 4 行灰色的原料名是自动的：哪一列是什么料，在【物料档案】原料的「生产日报第几列」里定；用过的列号别改给别的料，不用的料也别删，'
           '要加新料就给它一个没用过的列号（1～12）。不要在这张表里插列。产量是 0 但投了料也照登（成本留到下个月一起算）。'
           '期初盘点不在这里登（填物料档案的期初数量）。月底实盘在【月末盘点】登。灰底是示例行。')
    W = {C['日期']: 11, C['产品']: 11, C['产量']: 8, C['废料']: 9, C['备注']: 18, S['成品吨']: 10, S['投料合计']: 10, S['校验']: 40}
    for c in PROD_SLOT_COLS:
        W[c] = 10
    fm = {C['日期']: DATE, C['产量']: QTY, C['废料']: QTY, S['成品吨']: QTY, S['投料合计']: QTY}
    for c in PROD_SLOT_COLS:
        fm[c] = QTY
    al = {C['备注']: AL, C['产品']: AL}
    cols_all = [c for c, _ in heads] + PROD_SLOT_COLS
    rows = ctx.get('prod', [])
    widths(ws, W)
    title(ws, ws.title, C['备注'], C_IN, tip)
    header(ws, HDR, heads, C_IN, height=48)
    for j, c in enumerate(PROD_SLOT_COLS, 1):
        f = (f'=IFERROR(IF(N(INDEX(槽_料号,{j}))>0,INDEX(槽_名称,{j})&CHAR(10)&"（吨）","原料{j}（物料档案里还没指定）"),"原料{j}")')
        put(ws, f'{c}{HDR}', f, F_HDR, fill(H_SLOT), align=ACW)
    style_rows(ws, R0, R0 + N_PROD - 1, sorted(cols_all, key=CI), auto=tuple(S.values()), fmts=fm, aligns=al)
    for col in S.values():
        ws[f'{col}{HDR}'].fill = fill(H_AUTO)
    ws.freeze_panes = f'C{R0}'
    ws.auto_filter.ref = f'A{HDR}:{PROD_LAST}{R0 + N_PROD - 1}'
    last = R0 + N_PROD + 2000
    _label(ws, 'A3:B3', '筛选后合计 →')
    for c in [C['产量'], C['废料']] + PROD_SLOT_COLS + [S['成品吨'], S['投料合计']]:
        _sub(ws, c, last, QTY)
    _nofx(ws, 'A', S['校验'], f'{S["校验"]}3', last)
    _write_dicts(ws, rows, C)
    end = 'INDEX(产_末行,{n})=0'
    _show(ws, N_PROD, {
        S['成品吨']: f'IF({end},"",INDEX(产_成品吨,{{n}}))',
        S['投料合计']: f'IF({end},"",INDEX(产_投料合计,{{n}}))',
        S['校验']: _chk('产'),
    }, {S['成品吨']: QTY, S['投料合计']: QTY}, {S['成品吨']: AR, S['投料合计']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_PROD - 1)
    rng = rng_of(N_PROD)
    dv_day(ws, rng(C['日期']))
    dv_names(ws, rng(C['产品']), '产品列表', '空着＝默认产品（十水焦钠）；以后加了产品从这里选')
    dv_dec(ws, rng(C['产量']), '产量', prompt='成品入库多少袋（1 袋几吨在物料档案产品里填）')
    dv_dec(ws, rng(C['废料']), '废料', prompt='废料（二次过滤）多少袋，按回收料入库')
    dv_dec(ws, f'{PROD_SLOT_COLS[0]}{R0}:{PROD_SLOT_COLS[-1]}{R0 + N_PROD - 1}', '投料',
           prompt='这一列是什么料看第 4 行（在【物料档案】原料的「生产日报第几列」里定）；填投了多少吨')
    _tail(ws, PROD_LAST, '生产末行', N_PROD)


# ───────────────────────── 工资表 ─────────────────────────
def build_wage(ws, ctx):
    C, S = WAGE_COLS, WAGE_SHOW
    sub = {'月份': '1～12', '部门': '空＝人员表的', '应发合计': '空＝前四项相加', '扣社保': '个人部分',
           '扣借支': '预支过的，只扣', '实发': '空＝应发－扣款'}
    heads = [(C[k], _hd(k, WAGE_REQ, sub.get(k, ''))) for k in C] + [
        (S['应发'], '应发\n（自动）'), (S['实发算'], '实发\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 每月每人一行，月份填 1～12（几月的工资）。应发合计空着＝基本工资＋计件提成＋加班＋补贴，实发空着＝应发－扣社保－扣借支－其他扣款。'
           '月底按部门自动计提：生产→生产成本（生产工资）、车间管理→制造费用、管理→管理费用、销售→销售费用。'
           '发工资在【资金台帐】记「员工工费」；预支工资在资金台帐记「员工借支」，这里的扣借支只是扣回来，不另记账。'
           '发了工资一定要登这个月的工资表，不然成本少算（数据校验会提醒）。临时工、装卸工直接付的记资金台帐「劳务工资」，不登这里。灰底是示例行。')
    W = {C['月份']: 7, C['姓名']: 9, C['部门']: 9, C['岗位']: 9, C['出勤天数']: 8, C['基本工资']: 11, C['计件提成']: 10,
         C['加班']: 9, C['补贴']: 9, C['应发合计']: 11, C['扣社保']: 9, C['扣借支']: 9, C['其他扣款']: 9, C['实发']: 11,
         C['备注']: 16, S['应发']: 12, S['实发算']: 12, S['校验']: 40}
    money = ('基本工资', '计件提成', '加班', '补贴', '应发合计', '扣社保', '扣借支', '其他扣款', '实发')
    fm = {C[k]: MONEY for k in money}
    fm.update({C['月份']: MON_FMT, C['出勤天数']: '0.#'})
    al = {C[k]: AR for k in money}
    al[C['备注']] = AL
    rows = ctx.get('wages', [])
    _frame(ws, WAGE_LAST, heads, C_IN, tip, W, N_WAGE, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=C['备注'])
    last = R0 + N_WAGE + 2000
    _label(ws, 'A3:E3', '筛选后看得见的行合计 →')
    for k in money:
        _sub(ws, C[k], last)
    _sub(ws, S['应发'], last)
    _sub(ws, S['实发算'], last)
    _nofx(ws, 'A', S['校验'], f'{S["校验"]}3', last)
    _write_dicts(ws, rows, C)
    end = 'INDEX(工_末行,{n})=0'
    _show(ws, N_WAGE, {
        S['应发']: f'IF({end},"",INDEX(工_应发,{{n}}))',
        S['实发算']: f'IF({end},"",INDEX(工_实发,{{n}}))',
        S['校验']: _chk('工'),
    }, {S['应发']: MONEY, S['实发算']: MONEY}, {S['应发']: AR, S['实发算']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_WAGE - 1)
    rng = rng_of(N_WAGE)
    dv_whole(ws, rng(C['月份']), 1, 12, '月份', prompt='几月的工资，填 1～12')
    dv_names(ws, rng(C['姓名']), '人员列表', '从【基础资料】⑤人员里选；新员工先去人员里加一行')
    dv_fixed(ws, rng(C['部门']), DEPTS, prompt='空着＝人员表里的部门。生产→生产成本、车间管理→制造费用、管理→管理费用、销售→销售费用')
    dv_dec(ws, rng(C['出勤天数']), '出勤天数', '0', '31')
    for kk in money:
        dv_dec(ws, rng(C[kk]), kk)
    _tail(ws, WAGE_LAST, '工资末行', N_WAGE)


# ───────────────────────── 发票登记 ─────────────────────────
def build_inv(ws, ctx):
    C, S = INV_COLS, INV_SHOW
    req = globals().get('INV_REQ', ())
    sub = {'方向': '空＝按单位类型', '税额': '空＝按票种算', '对应单号': '进货日期车牌、合同号', '垫付编号': '补垫付、报销的票'}
    heads = [(C[k], _hd(k, req, sub.get(k, ''))) for k in C] + [(S['不含税'], '不含税\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 收到的发票一张一行（开给客户的票在【销售台账】填发票号就行，不用再登）。补到老板垫付、员工报销那一笔的票：垫付编号填资金台帐里那一笔的编号，'
           '进项税在登票这个月抵扣，原来那一行不用改。税额空着＝按票种的税率自动算。红字发票价税合计填负数。'
           '发票号码这一列是文本格式，直接输。灰底是示例行。')
    W = {C['日期']: 11, C['方向']: 8, C['往来单位']: 14, C['发票号码']: 22, C['票种']: 9, C['价税合计']: 13, C['税额']: 11,
         C['对应单号']: 16, C['垫付编号']: 12, C['备注']: 22, S['不含税']: 13, S['校验']: 40}
    fm = {C['日期']: DATE, C['价税合计']: MONEY, C['税额']: MONEY, C['发票号码']: '@', C['对应单号']: '@', C['垫付编号']: '@',
          S['不含税']: MONEY}
    al = {C[k]: AL for k in ('往来单位', '发票号码', '对应单号', '垫付编号', '备注')}
    al.update({C['价税合计']: AR, C['税额']: AR})
    rows = ctx.get('invoices', [])
    _frame(ws, INV_LAST, heads, C_IN, tip, W, N_INV, fm, al, f'D{R0}', auto=tuple(S.values()), title_col=C['备注'])
    last = R0 + N_INV + 2000
    _label(ws, 'A3:E3', '筛选后看得见的行合计 →')
    _sub(ws, C['价税合计'], last)
    _sub(ws, C['税额'], last)
    _nofx(ws, C['价税合计'], S['校验'], f'{C["对应单号"]}3', last, f'{C["备注"]}3')
    _sub(ws, S['不含税'], last)
    _write_dicts(ws, rows, C)
    end = 'INDEX(票_末行,{n})=0'
    _show(ws, N_INV, {
        S['不含税']: f'IF({end},"",INDEX(票_不含税,{{n}}))',
        S['校验']: _chk('票'),
    }, {S['不含税']: MONEY}, {S['不含税']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_INV - 1)
    rng = rng_of(N_INV)
    dv_day(ws, rng(C['日期']))
    dv_fixed(ws, rng(C['方向']), INV_DIRS, prompt='收到＝供应商、物流开给我们的；开出＝我们开给客户的（一般在销售台账填号就行）。空着＝按往来单位类型')
    dv_names(ws, rng(C['往来单位']), '往来单位列表', '从【往来单位】里选；菜市场这种零星的可以直接打字', stop=False)
    dv_names(ws, rng(C['票种']), '票种列表', '专票13%、专票9%、普票…；专票的税额才抵进项')
    dv_dec(ws, rng(C['价税合计']), '价税合计', prompt='发票上的价税合计（含税）；红字发票填负数')
    dv_dec(ws, rng(C['税额']), '税额', prompt='空着＝价税合计÷(1＋税率)×税率；跟票面差几分的照票面填')
    dv_tip(ws, rng(C['垫付编号']), '补到老板垫付、报销那一笔的票：填资金台帐里那一笔的编号（如 202608-X01），补票跟进里就消掉了')
    _tail(ws, INV_LAST, '发票末行', N_INV)


# ───────────────────────── 月末盘点 ─────────────────────────
def build_cnt(ws, ctx):
    C, S = CNT_COLS, CNT_SHOW
    sub = {'月份': '1～12', '物料': '原料或产品', '实盘数量': '按物料的单位'}
    heads = [(C[k], _hd(k, CNT_REQ, sub.get(k, ''))) for k in C] + [
        (S['账面数量'], '账面数量\n（自动，盘点前）'), (S['差异'], '差异\n（实盘－账面）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 月底盘点后，每种料、每种产品一行：月份、物料、实盘数量（原料按吨，产品按吨）。账面数量、差异自动出；'
           '差异按当月单价自动出盘点凭证（原料盘亏盘盈进制造费用，产品进营业外）。没盘的不用登；同一个月同一种料只登一行（重复的只认第一行）。'
           '不用盘点的辅料、吨袋（进货即耗用的）不用登。')
    W = {C['月份']: 7, C['物料']: 20, C['实盘数量']: 12, C['备注']: 24, S['账面数量']: 13, S['差异']: 12, S['校验']: 44}
    fm = {C['月份']: MON_FMT, C['实盘数量']: QTY, S['账面数量']: QTY, S['差异']: QTY}
    al = {C['物料']: AL, C['备注']: AL, C['实盘数量']: AR}
    rows = ctx.get('counts', [])
    _frame(ws, CNT_LAST, heads, C_IN, tip, W, N_CNT, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=CNT_LAST)
    last = R0 + N_CNT + 2000
    _nofx(ws, C['物料'], S['校验'], f'{C["备注"]}3', last, f'{S["校验"]}3')
    _write_dicts(ws, rows, C)
    ok = ('AND(INDEX(盘_有效,{n})=1,INDEX(盘_重复,{n})<>1,INDEX(盘_月,{n})>=1,INDEX(盘_月,{n})<=12)')
    m, k, p = 'INDEX(盘_月,{n})', 'INDEX(盘_料号,{n})', 'INDEX(盘_品号,{n})'
    book = (f'IF({ok},IF({k}>0,INDEX(料_末量,{k},{m})-INDEX(料_盘量,{k},{m}),'
            f'IF({p}>0,INDEX(品_末量,{p},{m})-INDEX(品_盘量,{p},{m}),"")),"")')
    e = f'INDEX(${S["账面数量"]}${HDR}:${S["账面数量"]}${REF_END},{{n}}+1)'
    _show(ws, N_CNT, {
        S['账面数量']: book,
        S['差异']: f'IF(ISNUMBER({e}),ROUND(N(INDEX(盘_实盘,{{n}}))-{e},3),"")',
        S['校验']: _chk('盘'),
    }, {S['账面数量']: QTY, S['差异']: QTY}, {S['账面数量']: AR, S['差异']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_CNT - 1)
    rng = rng_of(N_CNT)
    dv_whole(ws, rng(C['月份']), 1, 12, '月份', prompt='盘的是几月底，填 1～12')
    dv_names(ws, rng(C['物料']), '物料产品列表', '从【物料档案】里选原料、回收料或产品')
    dv_dec(ws, rng(C['实盘数量']), '实盘数量', '0', prompt='盘出来的数量（原料按吨、产品按吨；跟物料档案的单位一样）')
    _tail(ws, CNT_LAST, '盘点末行', N_CNT)


# ───────────────────────── 手工凭证（会计组） ─────────────────────────
def build_mj(ws, ctx):
    C, S = MJ_COLS, MJ_SHOW
    sub = {'凭证组号': '一张凭证一个号', '科目': '「编码 名称」或编码', '借方': '', '贷方': ''}
    heads = [(C[k], _hd(k, MJ_REQ, sub.get(k, ''))) for k in C if k != '往来明细'] + [
        (C['往来明细'], '往来/明细\n（应收应付填单位简称）'), (S['校验'], '这一行的问题\n（自动）')]
    heads.sort(key=lambda h: CI(h[0]))
    tip = ('★ 会计用：自动凭证以外的调账。同一张凭证同一个组号（一行一条分录），同组日期要一样，借贷合计相等才记账。'
           '科目从下拉选（「编码 名称」），也可以只输编码；应收账款、应付账款的往来/明细填单位简称（要在【往来单位】里）。'
           '不能用来调库存数量（盘点请用【月末盘点】），不能记现金、银行、微信、承兑（请记资金台帐）；记借款、垫付类科目会提醒。')
    W = {C['日期']: 11, C['凭证组号']: 9, C['摘要']: 24, C['科目']: 22, C['往来明细']: 14, C['借方']: 13, C['贷方']: 13,
         C['备注']: 18, S['校验']: 44}
    fm = {C['日期']: DATE, C['借方']: MONEY, C['贷方']: MONEY, C['科目']: '@'}
    al = {C[k]: AL for k in ('摘要', '科目', '往来明细', '备注')}
    al.update({C['借方']: AR, C['贷方']: AR})
    rows = ctx.get('manual', [])
    _frame(ws, MJ_LAST, heads, C_ACC, tip, W, N_MJ, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=MJ_LAST)
    last = R0 + N_MJ + 2000
    _label(ws, 'A3:E3', '筛选后看得见的行合计 →')
    _sub(ws, C['借方'], last)
    _sub(ws, C['贷方'], last)
    put(ws, f'{C["备注"]}3', f'=IF(ROUND({C["借方"]}3-{C["贷方"]}3,2)=0,"借贷平","借贷差 "&TEXT({C["借方"]}3-{C["贷方"]}3,"#,##0.00"))',
        F_AUTOB, FILL_AUTO, align=AC)
    _nofx(ws, C['科目'], S['校验'], f'{S["校验"]}3', last)
    _write_dicts(ws, rows, C)
    _show(ws, N_MJ, {S['校验']: _chk('手')}, {}, {S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_MJ - 1)
    rng = rng_of(N_MJ)
    dv_day(ws, rng(C['日期']))
    dv_tip(ws, rng(C['凭证组号']), '同一张凭证的几行填同一个号（1、2、3…或 J01），同组日期要一样，借贷相等才记账')
    dv_names(ws, rng(C['科目']), '科目列表', '从下拉选「编码 名称」，也可以只输编码。不能记现金、银行、微信、承兑、原材料、库存商品、生产成本', stop=False,
             error='不在会计科目里（科目先到【会计科目】加明细）')
    dv_names(ws, rng(C['往来明细']), '往来单位列表', '应收账款、应付账款填单位简称（要在【往来单位】里）；别的科目可以写明细说明', stop=False)
    for kk in ('借方', '贷方'):
        dv_dec(ws, rng(C[kk]), kk)
    _tail(ws, MJ_LAST, '手工末行', N_MJ)


# ───────────────────────── 基础资料 ─────────────────────────
PAR_NOTES = {
    '公司名称': '凭证、报表、对账单的抬头',
    '建账日': '从哪天开始用这本账（必须是某月 1 号）：账户余额、期初应收应付、库存都是前一天晚上的数，只录这一天及以后的单子。一本账管一个自然年，第二年用【年末结转】开新本',
    '纳税人类型': '一般纳税人：专票的进项可以抵；小规模：进项都进成本，销项按征收率',
    '销项税率': '一般纳税人卖货开票的税率（13%）',
    '征收率': '小规模纳税人的征收率',
    '销售含税': '销售台账的单价、金额是不是含税价（是＝含税）',
    '进货含税': '进货物流登记的单价、货款是不是含税价（是＝含税）',
    '运费税点': '运费开票加收的税点（7%）；运费票种是无票的不加',
    '销货运费记入': '我方出的销货运费记到哪：主营业务成本（算进成本）或销售费用',
    '投产月份': '写成 2026-07。建厂费用（环评、认证、装修…）从这个月开始摊；这个月以后付的认证、环评直接进费用',
    '摊销月数': '建厂费用、长期待摊在收支类别里没填分摊月数的，默认摊几个月',
    '默认产品': '生产日报、销售台账产品空着＝这个产品',
    '默认回收料': '废料（二次过滤）入库到哪种回收料（产品档案没填废料归入时用）',
    '回款天数': '客户欠款超过多少天没回款要提醒',
    '对账工作日': '发货（签收）后几个工作日对账（销售台账的结算情况用）',
    '付款天数': '对账后几天内付款，超过算逾期',
    '会计主管': '记账凭证下面签字处印的名字（可空）', '记账': '同上', '出纳': '同上', '审核': '同上', '制单': '同上',
    '凭证号打印': '记账凭证上要不要印凭证号',
}
BLOCK_HEADS = {
    '账户': {'名称': '账户名称★', '类型': '类型', '期初余额': '期初余额\n（建账日前一天晚上）', '人员': '人员\n（垫付人、拿备用金的人）', '备注': '备注'},
    '类别': {'名称': '收支类别★', '归类': '归类\n（算到哪）', '收钱科目': '收钱记到\n（科目）', '付钱科目': '付钱记到\n（科目）',
           '现流收': '现金流量\n（收，行次）', '现流付': '现金流量\n（付，行次）', '默认票种': '默认票种', '不得抵扣': '进项\n不得抵扣',
           '分摊月数': '分摊月数\n（>1 先挂账按月摊）', '摊进': '摊进\n（制造/管理）', '投产后科目': '投产后记到\n（建厂费用类）',
           '默认往来': '默认往来单位', '说明': '说明'},
    '报销': {'名称': '报销项目★', '科目': '记到\n（科目）', '默认票种': '默认票种', '不得抵扣': '进项\n不得抵扣', '现流': '现金流量\n（行次）', '说明': '说明'},
    '人员': {'姓名': '姓名★', '部门': '部门', '备注': '备注'},
    '票种': {'名称': '票种★', '可抵税率': '可抵税率'},
    '借期': {'出借人': '出借人★', '日期': '日期★\n（打款日）', '金额': '金额★', '性质': '性质★\n（借款/投资/垫付）', '已付利息': '建账前\n已付利息',
           '说明': '说明'},
    '期初': {'科目': '科目★', '明细': '明细\n（单位、项目）', '借方': '借方余额', '贷方': '贷方余额', '剩余月数': '剩余摊销月数\n（1801、1123）',
           '摊进': '摊进\n（制造/管理）', '说明': '说明'},
}
BLOCK_NOTES = {
    '账户': ('收钱、付钱的地方。类型：现金、银行、微信支付宝是真钱（进现金流量表、可用资金）；票据＝手上的承兑汇票；'
           '垫付＝老板（或合伙人）个人替公司付钱的账，人员填垫付人（【往来单位】里要有这个出借人），期初不填（建账前垫的填⑦）；'
           '备用金＝发给员工的备用金，人员填拿钱的人。期初余额填建账日前一天晚上的余额（按网银核对）。这一块最多 {cap} 个。'),
    '类别': ('资金台帐的收支类别下拉，名字照老板流水账的项目（原名原顺序）。收钱/付钱记到：科目（下拉「编码 名称」），也可以填 '
           '@往来（按往来单位记应收应付）、@借款、@利息、@分红、@报销（按报销项目）、@垫付人、@备用金、@投资、@停用（选了报错）。'
           '现金流量行次：1 销售收到 2 收其他经营 3 购买原材料 4 职工薪酬 5 税费 6 付其他经营 10 处置长期资产 12 购建长期资产 '
           '14 取得借款 15 吸收投资 16 偿还借款 17 偿还利息 18 分配利润，0＝不进。分摊月数大于 1：付款先挂长期待摊（1801）或预付（1123），'
           '按月摊；摊进：建厂费用摊到制造费用还是管理费用；投产后科目：投产月以后付的直接记这个科目。这一块最多 {cap} 个。'),
    '报销': ('员工报销的细项（材料配件、加油、日常用品…），报销统计按报销人×报销项目出。科目决定记到哪；默认票种：报销时票种空着就用它；'
           '餐费、招待这类进项不能抵的「不得抵扣」选是。这一块最多 {cap} 个。'),
    '人员': ('报销人、工资表姓名、拿备用金的人的下拉。部门决定工资记到哪：生产→生产成本（生产工资）、车间管理→制造费用、'
           '管理→管理费用、销售→销售费用。这一块最多 {cap} 个。'),
    '票种': ('资金台帐、进货、销售、发票登记的票种下拉。可抵税率：专票填税率（一般纳税人按它抵进项），普票、无票、待补票填 0。'
           '「待补票」是老板垫付、报销先没票的，补到票在【发票登记】登。这一块最多 {cap} 个。'),
    '借期': ('建账日以前借进来、投进来、老板垫付还没还的钱，一笔一行（带打款日，按天计息）。出借人要在【往来单位】建成「出借人」并填年利率。'
           '性质：借款（银行的记短期借款，个人、合伙人的记其他应付款）、投资（记实收资本，按投资年利率算「利息」）、垫付（老板垫的，记垫付款）。'
           '建账前已经付过的利息填在「建账前已付利息」。这些期初只填这里，不要再在②或⑧填。这一块最多 {cap} 笔。'),
    '期初': ('其他科目的期初余额（押金、建厂费用没摊完的、欠的工资、税、跨月费用……），一个科目（明细）一行，借方、贷方只填一个。'
           '现金银行（②）、存货（物料档案）、应收应付（往来单位）、借款投资垫付利息（⑦）、固定资产（卡片）不填这里。'
           '长期待摊（1801）、预付（1123）没摊完的填剩余摊销月数和摊进。未分配利润不用填，自动倒挤。这一块最多 {cap} 行。'),
}


def build_base(ws, ctx):
    NOTE = 'O'
    widths(ws, {'A': 16, 'B': 12, 'C': 13, 'D': 12, 'E': 10, 'F': 10, 'G': 10, 'H': 8, 'I': 10, 'J': 9, 'K': 11, 'L': 12,
                'M': 30, 'N': 2, NOTE: 70})
    title(ws, '基 础 资 料', 'M', C_MASTER,
          '★ 一次填好、偶尔改：① 参数；② 账户；③ 收支类别（资金台帐的下拉，决定记到哪个科目、现金流量哪一行）；④ 报销项目；⑤ 人员；'
          '⑥ 票种；⑦ 建账前的借款、投资、垫付（一笔一行，计息靠它）；⑧ 其他期初余额。每块上下排着，在块里插整行、删整行只动这一块；'
          '块的标题（② 账户…）别改。名称不要重复；改了名字，录入表里的旧名字要用「查找替换」一起改。黄格是要填的，灰底是示例（正式用之前删掉）。')
    home_link(ws, f'{NOTE}1')
    pv = ctx.get('params', {})
    section(ws, BASE_PAR_R0 - 1, 'A', 'D', '① 参数', C_MASTER)
    for k, lbl, dflt in PARAMS:
        r = PAR_ROW[k]
        put(ws, f'A{r}', lbl, F_KPI_L, F_LBL, align=ALW)
        ws.merge_cells(f'A{r}:B{r}')
        v = pv.get(k, dflt)
        if k == '建账日':
            v = _v(v)
        fmt = {'建账日': DATE, '销项税率': '0.0%', '征收率': '0.0%', '运费税点': '0.0%', '投产月份': '@'}.get(k)
        put(ws, f'{BASE_PAR_VAL}{r}', v, F_IN, FILL_IN, fmt, AC)
        put(ws, f'{BASE_PAR_NOTE}{r}', PAR_NOTES.get(k, ''), F_NOTE, align=AL, border=False)
        ws.merge_cells(f'{BASE_PAR_NOTE}{r}:M{r}')
        ws.row_dimensions[r].height = 30
    P = lambda k: f'{BASE_PAR_VAL}{PAR_ROW[k]}'
    dv_day(ws, P('建账日'), '某月 1 号，例如 2026-8-1')
    dv_fixed(ws, P('纳税人类型'), TAXPAYER)
    for k in ('销项税率', '征收率', '运费税点'):
        dv_dec(ws, P(k), k, '0', '1', '（13% 就填 13%）')
    for k in ('销售含税', '进货含税', '凭证号打印'):
        dv_fixed(ws, P(k), YESNO)
    dv_fixed(ws, P('销货运费记入'), FREIGHT_TO)
    dv_tip(ws, P('投产月份'), '写成 2026-07 这样（年-月）')
    dv_whole(ws, P('摊销月数'), 1, 240, '摊销月数')
    dv_names(ws, P('默认产品'), '产品列表', '【物料档案】② 产品里的名字', stop=False)
    dv_names(ws, P('默认回收料'), '原料列表', '【物料档案】① 原料里类别是回收料的那一种', stop=False)
    for k in ('回款天数', '对账工作日', '付款天数'):
        dv_whole(ws, P(k), 0, 999, k)
    for k in ('会计主管', '记账', '出纳', '审核', '制单'):
        dv_len(ws, P(k), 10)
    data = {'账户': ctx.get('accounts', []), '类别': ctx.get('cats', []), '报销': ctx.get('reimb', []),
            '人员': ctx.get('persons', []), '票种': ctx.get('tickets', []), '借期': ctx.get('loans0', []),
            '期初': ctx.get('opens', [])}
    note_key = {'账户': '备注', '类别': '说明', '报销': '说明', '人员': '备注', '票种': None, '借期': '说明', '期初': '说明'}
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        cols = BASE_COLS[key]
        c2 = max(cols.values(), key=CI)
        section(ws, tr, 'A', c2, ttl, C_MASTER)
        assert ws[f'A{tr}'].value == ttl
        put(ws, f'{NOTE}{tr}', BLOCK_NOTES[key].format(cap=cap), F_NOTE, align=Alignment(horizontal='left', vertical='top',
                                                                                          wrap_text=True), border=False)
        ws.merge_cells(f'{NOTE}{tr}:{NOTE}{tr + min(cap, 9)}')
        header(ws, hr, [(cols[k], BLOCK_HEADS[key][k]) for k in cols], C_MASTER, height=44)
        r0 = hr + 1
        fm, al = {}, {}
        for k, c in cols.items():
            if k in ('名称', '姓名', '出借人', '说明', '备注', '明细', '默认往来'):
                al[c] = AL
            if k in ('收钱科目', '付钱科目', '投产后科目', '科目'):
                fm[c] = '@'
                al[c] = AL
            if k in ('期初余额', '金额', '已付利息', '借方', '贷方'):
                fm[c] = MONEY
                al[c] = AR
            if k == '日期':
                fm[c] = DATE
            if k == '可抵税率':
                fm[c] = '0%'
        style_rows(ws, r0, r0 + cap - 1, list(cols.values()), fmts=fm, aligns=al, fills={})
        _write_tuples(ws, data[key][:cap], cols, r0, note_key[key])
        R = lambda k: f'{cols[k]}{r0}:{cols[k]}{r0 + cap - 1}'
        if key in ('账户', '类别', '报销', '人员', '票种'):                     # A 列是名称（借期、期初的 A 列另有下拉）
            dv_len(ws, f'A{r0}:A{r0 + cap - 1}', 30)
        if key == '账户':
            dv_fixed(ws, R('类型'), ACC_TYPES, prompt='现金、银行、微信支付宝是真钱；票据＝承兑汇票；垫付＝老板个人替公司付钱的账；备用金＝发给员工的。空着＝银行')
            dv_dec(ws, R('期初余额'), '期初余额', prompt='建账日前一天晚上的余额（按网银核对）。垫付账户不填（建账前垫的填⑦）')
            dv_names(ws, R('人员'), '人员列表', '垫付账户填垫付人（往来单位里的出借人，如老板）；备用金账户填拿钱的人', stop=False)
        elif key == '类别':
            dv_fixed(ws, R('归类'), KINDS, prompt='资金月报、「算到哪」用的大类')
            for k in ('收钱科目', '付钱科目', '投产后科目'):
                dv_names(ws, R(k), '科目列表', '下拉选「编码 名称」；也可以填 @往来 @借款 @利息 @分红 @报销 @垫付人 @备用金 @投资 @停用（见右边说明）',
                         stop=False, error='不在会计科目里，也不是 @往来 这类记号')
            for k in ('现流收', '现流付'):
                dv_whole(ws, R(k), 0, 20, '现金流量行次', prompt='1 销售收到 2 收其他经营 3 购买原材料 4 职工薪酬 5 税费 6 付其他经营 '
                                                           '10 处置长期资产 12 购建长期资产 14 取得借款 15 吸收投资 16 偿还借款 17 偿还利息 18 分配利润；0＝不进')
            dv_names(ws, R('默认票种'), '票种列表', '这个类别的钱票种空着时用它')
            dv_fixed(ws, R('不得抵扣'), YESNO, prompt='招待、伙食、福利这类进项不能抵的选是')
            dv_whole(ws, R('分摊月数'), 0, 240, '分摊月数', prompt='一次付几个月的（年付房租、建厂费用）填摊几个月；空或 1＝当月就算')
            dv_fixed(ws, R('摊进'), SPREAD_TO, prompt='建厂费用（长期待摊 1801）摊到制造费用还是管理费用')
            dv_names(ws, R('默认往来'), '往来单位列表', '选了这个类别、往来单位空着时用它（如微众贷款利息→微众银行）', stop=False)
        elif key == '报销':
            dv_names(ws, R('科目'), '科目列表', '下拉选「编码 名称」', stop=False, error='不在会计科目里')
            dv_names(ws, R('默认票种'), '票种列表', '报销时票种空着用它')
            dv_fixed(ws, R('不得抵扣'), YESNO, prompt='餐费、招待这类进项不能抵的选是')
            dv_whole(ws, R('现流'), 0, 20, '现金流量行次', prompt='3 购买原材料 4 职工薪酬 6 付其他经营…（见③右边说明）')
        elif key == '人员':
            dv_fixed(ws, R('部门'), DEPTS, prompt='决定工资记到哪：生产、车间管理、管理、销售')
        elif key == '票种':
            dv_dec(ws, R('可抵税率'), '可抵税率', '0', '1', prompt='专票填税率（13%），普票、无票、待补票填 0')
        elif key == '借期':
            dv_names(ws, R('出借人'), '出借人列表', '【往来单位】里类型是出借人的（银行、个人、合伙人、老板）')
            dv_day(ws, R('日期'), '钱打过来（或老板垫付）的那一天，要早于建账日')
            dv_dec(ws, R('金额'), '金额', prompt='到建账日还没还的本金')
            dv_fixed(ws, R('性质'), LO_KINDS, prompt='借款：借来的钱；投资：合伙人投的钱（记实收资本）；垫付：老板个人替公司付的、还没还的')
            dv_dec(ws, R('已付利息'), '建账前已付利息', prompt='建账前已经付过的这一笔的利息（从应计利息里扣）')
        elif key == '期初':
            dv_names(ws, R('科目'), '科目列表', '下拉选「编码 名称」。现金银行、存货、应收应付、借款投资、固定资产不填这里', stop=False,
                     error='不在会计科目里')
            for k in ('借方', '贷方'):
                dv_dec(ws, R(k), k)
            dv_whole(ws, R('剩余月数'), 0, 240, '剩余摊销月数', prompt='长期待摊（1801）、预付（1123）建账时还要摊几个月')
            dv_fixed(ws, R('摊进'), SPREAD_TO, prompt='1801 摊到制造费用还是管理费用')
    ws.freeze_panes = 'A3'
    print_setup(ws, None, landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'


# ───────────────────────── 往来单位 ─────────────────────────
def build_unit(ws, ctx):
    C, S = UNIT_COLS, UNIT_SHOW
    req = ('简称', '类型')
    sub = {'别名': '几个用、隔开', '期初应收': '客户欠我们', '期初应付': '我们欠对方', '期初未开票': '可不填', '期初欠票': '可不填',
           '开票': '空＝要', '默认票种': '个人的选无票', '出借人类别': '出借人填', '借款年利率': '月息1分＝12%，8厘＝9.6%',
           '投资年利率': '合伙人投资', '计息天数': '空＝银行360，个人365', '到期日': '借款'}
    heads = [(C[k], _hd(k, req, sub.get(k, ''))) for k in C] + [(S['现在'], '现在\n（自动，到截止日）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 客户、供应商、物流、出借人都在这里建档，一家一行。简称是下拉和报表里用的名字；全称、其他叫法（几个用、隔开）也登上，录入时写哪个都认。'
           '期初应收、期初应付填建账日前一天晚上的余额（预收、预付填负数）；出借人不填期初，建账前借的钱填【基础资料】⑦。'
           '出借人（银行、个人、合伙人、老板）要填出借人类别、年利率（月息1分＝12%，8厘＝9.6%）；银行每月自动扣息的可以不填利率（付息直接进财务费用）。'
           '同一家两笔利率不一样的，建两个（农行-1、农行-2）。个人司机、个人料商默认票种选无票（税点就不加）。灰底是示例单位，正式用时连同各表的示例行一起删掉。')
    W = {C['简称']: 13, C['类型']: 9, C['全称']: 24, C['别名']: 14, C['联系人']: 8, C['电话']: 12, C['地址']: 18,
         C['期初应收']: 12, C['期初应付']: 12, C['期初未开票']: 11, C['期初欠票']: 11, C['开票']: 7, C['默认票种']: 9,
         C['出借人类别']: 9, C['借款年利率']: 11, C['投资年利率']: 10, C['计息天数']: 9, C['到期日']: 11, C['备注']: 26,
         S['现在']: 24, S['校验']: 40}
    money = ('期初应收', '期初应付', '期初未开票', '期初欠票')
    fm = {C[k]: MONEY for k in money}
    fm.update({C['电话']: '@', C['借款年利率']: '0.00%', C['投资年利率']: '0.00%', C['到期日']: DATE, C['计息天数']: '0'})
    al = {C[k]: AL for k in ('简称', '全称', '别名', '地址', '备注', '电话')}
    for k in money + ('借款年利率', '投资年利率'):
        al[C[k]] = AR
    rows = ctx.get('units', [])
    _frame(ws, UNIT_LAST, heads, C_MASTER, tip, W, N_UNIT, fm, al, f'B{R0}', auto=tuple(S.values()), title_col=C['备注'])
    _kpi(ws, 3, [('A', '期初应收合计', '=SUMIFS(位_期初应收,位_是客户,1)', MONEY),
                 ('D', '期初应付合计', '=SUMIFS(位_期初应付,位_是供应商,1)', MONEY),
                 ('G', '客户 / 供应商', '=SUM(位_是客户)&" 家 / "&SUM(位_是供应商)&" 家"', '@')])
    _nofx(ws, 'A', S['校验'], f'{C["期初欠票"]}3', R0 + N_UNIT + 1000, f'{C["备注"]}3')
    for i, u in enumerate(rows):
        r = R0 + i
        for k, c in C.items():
            v = _v(u.get(k))
            if v is not None:
                ws[f'{c}{r}'] = v
        if _is_ex(u.get('备注')):
            _grey(ws, r, 'A', C['备注'])
    yr, yp = 'INDEX(位_应收,{n})', 'INDEX(位_应付,{n})'
    isc, iss = 'INDEX(位_是客户,{n})=1', 'INDEX(位_是供应商,{n})=1'
    zr, zp = f'ABS({yr})<0.005', f'ABS({yp})<0.005'
    now = (f'IF(INDEX(位_有效,{{n}})<>1,"",IF(OR(NOT({isc}),{zr}),"",IF({yr}<0,"预收 ","应收 ")&TEXT(ABS({yr}),"#,##0.00"))'
           f'&IF(AND({isc},{iss},NOT({zr}),NOT({zp})),"；","")&IF(OR(NOT({iss}),{zp}),"",IF({yp}<0,"预付 ","应付 ")&TEXT(ABS({yp}),"#,##0.00")))')
    _show(ws, N_UNIT, {S['现在']: now, S['校验']: _chk1('位')}, {}, {S['现在']: AL, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_UNIT - 1)
    rng = rng_of(N_UNIT)
    dv_fixed(ws, rng(C['类型']), UNIT_TYPES, prompt='两边都有往来的选「客户和供应商」；借钱给公司的银行、个人、合伙人、老板选「出借人」')
    dv_fixed(ws, rng(C['开票']), ['要', '不要'], prompt='空着＝要；个人、不开票的选「不要」（不算未开票、欠票）')
    dv_names(ws, rng(C['默认票种']), '票种列表', '进货、运费、销售行票种空着时用它；个人司机、个人料商选无票')
    dv_fixed(ws, rng(C['出借人类别']), LENDER_KINDS, prompt='银行：借款记短期借款；个人、合伙人：记其他应付款-个人及合伙人借款')
    for k in ('借款年利率', '投资年利率'):
        dv_dec(ws, rng(C[k]), k, '0', '1', '（12% 就填 12%）', prompt='年利率：月息 1 分＝12%，8 厘＝9.6%。银行每月自动扣息的可以不填')
    dv_whole(ws, rng(C['计息天数']), 360, 366, '计息天数', prompt='一年按几天算利息：银行一般 360，个人 365；空着自动')
    dv_day(ws, rng(C['到期日']), '借款到期日（首页提前 30 天提醒）')
    for k in money:
        dv_dec(ws, rng(C[k]), k)
    dv_len(ws, rng(C['简称']), 20)
    _tail(ws, UNIT_LAST, '往来末行', N_UNIT)


# ───────────────────────── 物料档案 ─────────────────────────
def build_item(ws, ctx):
    C, P, SR, SP = RM_COLS, PRD_COLS, ITEM_SHOW_RM, ITEM_SHOW_PRD
    last = SR['校验']
    W = {'A': 20, 'B': 16, 'C': 9, 'D': 7, 'E': 10, 'F': 13, 'G': 10, 'H': 8, 'I': 10, 'J': 10, 'K': 10, 'L': 12, 'M': 11,
         'N': 10, 'O': 18, 'P': 12, 'Q': 14, 'R': 40}
    widths(ws, W)
    tip = ('★ ① 原料：原料、辅料、包装物、回收料（废料）都在这里，一种一行。「生产日报第几列」给生产日报里占一列的料编号（1～12，不能重复；'
           '用过的列号不要改给别的料，不用的料也别删）。辅料、包装物默认「进货即耗用」（买了当月就进成本，不管库存），想管库存的填「否」并给它一列。'
           '参考单价：没定价的进货按它暂估。期初数量、期初单价填建账日前一天的盘点数（有数量一定要有单价或金额）。'
           '② 产品：每袋吨数（1 袋＝1 吨填 1）、分摊系数（几种产品分制造费用的权重，空＝1）、废料归入（回收料的名字）。灰底是示例。')
    title(ws, ws.title, last, C_MASTER, tip)
    home_link(ws, f'{CL(CI(last) + 1)}1')
    ws.freeze_panes = 'B4'
    # ① 原料
    r_t, hr = ITEM_R0, ITEM_RM_HDR
    section(ws, r_t, 'A', last, f'① 原料（含辅料、包装物、回收料；最多 {N_RM} 种）', C_MASTER)
    sub = {'别名': '几个用、隔开', '类别': '原料/辅料/包装物/回收料', '单位': '吨/个/kg', '进货即耗用': '空：辅料包装物＝是',
           '计入明细': '进货即耗用记到', '生产日报列': '1～12，不能重复', '含量': '%', '参考单价': '待定价暂估用',
           '期初数量': '建账日前一天', '期初金额': '空＝数量×单价', '定额单价': '回收料入库价', '每袋吨数': '回收料'}
    hd = {'生产日报列': '生产日报第几列'}
    heads = [(C[k], _hd(hd.get(k, k), ('名称',), sub.get(k, ''))) for k in C] + [
        (SR['本年进'], '本年进\n（自动）'), (SR['现在库存'], '现在库存\n（自动，到截止日）'), (SR['校验'], '这一行的问题\n（自动）')]
    header(ws, hr, heads, C_MASTER, height=44)
    for col in SR.values():
        ws[f'{col}{hr}'].fill = fill(H_AUTO)
    r0 = hr + 1
    fm = {C['含量']: '0.0%', C['参考单价']: PRICE, C['期初数量']: QTY, C['期初单价']: PRICE, C['期初金额']: MONEY,
          C['定额单价']: PRICE, C['每袋吨数']: '0.###', C['计入明细']: '@', SR['本年进']: QTY, SR['现在库存']: QTY}
    al = {C[k]: AL for k in ('名称', '别名', '备注', '计入明细')}
    style_rows(ws, r0, r0 + N_RM - 1, [c for c, _ in heads], auto=tuple(SR.values()), fmts=fm, aligns=al)
    _write_tuples(ws, ctx.get('rms', [])[:N_RM], C, r0, '备注')
    A, rr = C['名称'], lambda c: f'${c}${r0}:${c}${r0 + N_RM - 1}'
    o = lambda c: own(c, hr)
    nm = f'TRIM({o("A")}&"")'
    kd = f'TRIM({o(C["类别"])}&"")'
    use = f'TRIM({o(C["进货即耗用"])}&"")'
    g = o(C['生产日报列'])
    q0, p0, a0 = o(C['期初数量']), o(C['期初单价']), o(C['期初金额'])
    is_use = f'OR({use}="是",AND({use}="",OR({kd}="辅料",{kd}="包装物")))'
    chk_rm = (f'IF(COUNTA({own_row("A", C["备注"], hr)})=0,"",IF({nm}="","✗ 没填名称（这一行不算）",'
              f'IF(COUNTIF({rr(A)},{esc(nm)})>1,"✗ 名称重复了",'
              f'IF(NOT(OR({kd}="",{kd}="原料",{kd}="辅料",{kd}="包装物",{kd}="回收料")),"✗ 类别要选原料、辅料、包装物、回收料",'
              f'IF(AND({g}<>"",OR(NOT(ISNUMBER({g})),N({g})<1,N({g})>{N_SLOT},N({g})<>INT(N({g})))),"✗ 生产日报第几列填 1～{N_SLOT} 的整数",'
              f'IF(AND(ISNUMBER({g}),COUNTIF({rr(C["生产日报列"])},{g})>1),"✗ 生产日报第 "&{g}&" 列给了两种料",'
              f'IF(AND(N({q0})<>0,NOT(ISNUMBER({p0})),NOT(ISNUMBER({a0}))),"✗ 期初有数量没单价、金额（不填成本会算错）",'
              f'IF(AND({is_use},ISNUMBER({g})),"⚠ 进货即耗用的料不用占生产日报的列（要在生产日报里领用，进货即耗用填否）",'
              f'IF(AND(NOT({is_use}),OR({kd}="辅料",{kd}="包装物"),NOT(ISNUMBER({g}))),"⚠ 不是进货即耗用、又没占生产日报的列：买进来出不去",'
              f'IF(AND(OR({kd}="",{kd}="原料"),NOT(ISNUMBER({g}))),"⚠ 原料没给生产日报的列：生产日报里领不了","✓"))))))))))')
    now = (f'IF(INDEX(料档_有效,{{n}})<>1,"",IF(INDEX(料档_耗用,{{n}})=1,"进货即耗用",'
           f'IF(P_截止月>=1,INDEX(料_末量,{{n}},P_截止月),"")))')
    _show(ws, N_RM, {
        SR['本年进']: f'IF(INDEX(料档_有效,{{n}})<>1,"",SUM(INDEX(料_进量,{{n}},0))+SUM(INDEX(料_回量,{{n}},0)))',
        SR['现在库存']: now,
        SR['校验']: chk_rm,
    }, {SR['本年进']: QTY, SR['现在库存']: QTY}, {SR['本年进']: AR, SR['现在库存']: AR, SR['校验']: AL}, r0=r0, hdr=hr)
    _check_cf(ws, SR['校验'], r0, r0 + N_RM - 1)
    R = rng_of(N_RM, r0)
    dv_fixed(ws, R(C['类别']), RM_KINDS, prompt='原料：生产日报里领用；辅料、包装物：默认买了就算成本；回收料：废料（二次过滤）')
    dv_fixed(ws, R(C['单位']), ['吨', 'kg', '个', '袋'], stop=False, error='不是常用单位（可以直接写，点「是」）')
    dv_fixed(ws, R(C['进货即耗用']), YESNO, prompt='是＝买了当月就进成本（不管库存）；空着：辅料、包装物＝是，原料＝否')
    dv_names(ws, R(C['计入明细']), '科目列表', '进货即耗用的记到哪个制造费用明细；空着：辅料→机物料消耗、包装物→包装物', stop=False,
             error='不在会计科目里')
    dv_whole(ws, R(C['生产日报列']), 1, N_SLOT, '生产日报第几列', prompt='这种料在生产日报原料第几列（1～12，不能重复）。用过的列号不要改给别的料')
    dv_dec(ws, R(C['含量']), '含量', '0', '100', prompt='83% 就填 83%')
    for k in ('参考单价', '期初单价', '定额单价'):
        dv_dec(ws, R(C[k]), k, prompt='每吨（每个单位）含税价' if k == '参考单价' else None)
    for k in ('期初数量', '期初金额'):
        dv_dec(ws, R(C[k]), k)
    dv_dec(ws, R(C['每袋吨数']), '每袋吨数', '0', '100', prompt='回收料（废料）1 袋几吨；空＝1')
    dv_len(ws, R(A), 30)
    # ② 产品
    r_t, hr = ITEM_PRD_R0, ITEM_PRD_HDR
    lastp = SP['校验']
    section(ws, r_t, 'A', lastp, f'② 产品（最多 {N_PRD} 种）', C_MASTER)
    sub = {'别名': '几个用、隔开', '包装单位': '袋', '每袋吨数': '1 袋＝1 吨填 1', '分摊系数': '空＝1', '废料归入': '回收料名',
           '成品含量': '%', '期初数量': '吨', '期初单价': '每吨', '期初金额': '空＝数量×单价', '期初在产': '在产品金额'}
    heads = [(P[k], _hd(k, ('名称',), sub.get(k, ''))) for k in P] + [
        (SP['本年产'], '本年产\n（吨，自动）'), (SP['现在库存'], '现在库存\n（吨，自动）'), (SP['校验'], '这一行的问题\n（自动）')]
    header(ws, hr, heads, C_MASTER, height=44)
    for col in SP.values():
        ws[f'{col}{hr}'].fill = fill(H_AUTO)
    r0 = hr + 1
    fm = {P['每袋吨数']: '0.###', P['分摊系数']: '0.###', P['折纯系数']: '0.####', P['成品含量']: '0.0%', P['期初数量']: QTY,
          P['期初单价']: PRICE, P['期初金额']: MONEY, P['期初在产']: MONEY, SP['本年产']: QTY, SP['现在库存']: QTY}
    al = {P[k]: AL for k in ('名称', '别名', '废料归入', '备注')}
    style_rows(ws, r0, r0 + N_PRD - 1, [c for c, _ in heads], auto=tuple(SP.values()), fmts=fm, aligns=al)
    _write_tuples(ws, ctx.get('prds', [])[:N_PRD], P, r0, '备注')
    rm_names = f'${A}${ITEM_RM_HDR + 1}:${A}${ITEM_RM_HDR + N_RM}'
    rm_alias = f'${C["别名"]}${ITEM_RM_HDR + 1}:${C["别名"]}${ITEM_RM_HDR + N_RM}'
    pr = lambda c: f'${c}${r0}:${c}${r0 + N_PRD - 1}'
    o = lambda c: own(c, hr)
    nm = f'TRIM({o("A")}&"")'
    bag = o(P['每袋吨数'])
    fq, fp, fa = o(P['期初数量']), o(P['期初单价']), o(P['期初金额'])
    wf = f'TRIM({o(P["废料归入"])}&"")'
    chk_p = (f'IF(COUNTA({own_row("A", P["备注"], hr)})=0,"",IF({nm}="","✗ 没填名称（这一行不算）",'
             f'IF(COUNTIF({pr("A")},{esc(nm)})>1,"✗ 名称重复了",IF(COUNTIF({rm_names},{esc(nm)})>0,"✗ 跟上面的原料重名了",'
             f'IF(AND({bag}<>"",OR(NOT(ISNUMBER({bag})),N({bag})<=0)),"✗ 每袋吨数要填正数（1 袋＝1 吨填 1）",'
             f'IF(AND(N({fq})<>0,NOT(ISNUMBER({fp})),NOT(ISNUMBER({fa}))),"✗ 期初有数量没单价、金额（不填利润会虚高）",'
             f'IF(AND({wf}<>"",COUNTIF({rm_names},{esc(wf)})=0,COUNTIF({rm_alias},"*"&{esc(wf)}&"*")=0),"⚠ 废料归入的料在上面原料里找不到",'
             f'IF({bag}="","⚠ 每袋吨数没填（按 1 吨算）","✓"))))))))')
    _show(ws, N_PRD, {
        SP['本年产']: f'IF(INDEX(品档_有效,{{n}})<>1,"",SUM(INDEX(品_完量,{{n}},0)))',
        SP['现在库存']: f'IF(INDEX(品档_有效,{{n}})<>1,"",IF(P_截止月>=1,INDEX(品_末量,{{n}},P_截止月),""))',
        SP['校验']: chk_p,
    }, {SP['本年产']: QTY, SP['现在库存']: QTY}, {SP['本年产']: AR, SP['现在库存']: AR, SP['校验']: AL}, r0=r0, hdr=hr)
    _check_cf(ws, SP['校验'], r0, r0 + N_PRD - 1)
    R = rng_of(N_PRD, r0)
    dv_fixed(ws, R(P['包装单位']), ['袋', '吨'], stop=False, error='不是常用单位（可以直接写，点「是」）')
    dv_dec(ws, R(P['每袋吨数']), '每袋吨数', '0', '100', prompt='生产日报按袋登，1 袋几吨（1000kg＝1）')
    dv_dec(ws, R(P['分摊系数']), '分摊系数', '0', '1000', prompt='几种产品分制造费用时每吨的权重，空＝1')
    dv_names(ws, R(P['废料归入']), '原料列表', '这个产品的废料（二次过滤）入到哪种回收料；空＝基础资料的默认回收料', stop=False)
    dv_dec(ws, R(P['成品含量']), '成品含量', '0', '100')
    for k in ('期初数量', '期初单价', '期初金额', '期初在产', '折纯系数'):
        dv_dec(ws, R(P[k]), k, prompt='建账日前一天盘点的成品吨数（有数量一定要填单价或金额）' if k == '期初数量' else None)
    dv_len(ws, R(P['名称']), 30)
    print_setup(ws, None, landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'


# ───────────────────────── 固定资产 ─────────────────────────
def build_fa(ws, ctx):
    C, S = FA_COLS, FA_SHOW
    req = ('名称', '取得日期', '原值', '折旧月数')
    sub = {'使用部门': '生产/管理/销售', '原值': '不含税', '残值率': '空＝5%', '折旧月数': '机器120 车48 电脑36',
           '建账前已提': '空＝自动算', '处置日期': '卖了、报废了填'}
    heads = [(C[k], _hd(k, req, sub.get(k, ''))) for k in C] + [
        (S['月折旧'], '月折旧\n（自动）'), (S['累计折旧'], '累计折旧\n（自动，到截止日）'), (S['净值'], '净值\n（自动）'),
        (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 设备、车辆一台一行（一张卡片）。原值填不含税价（跟资金台帐「机械设备费」那一笔的不含税额一致），取得的下个月开始提折旧，提满为止；'
           '卖了、报废了填处置日期。使用部门决定折旧进哪里：生产→制造费用，管理→管理费用，销售→销售费用。建账前已提空着＝按月数自动算。灰底是示例。')
    W = {C['编号']: 8, C['名称']: 16, C['类别']: 10, C['使用部门']: 9, C['取得日期']: 11, C['原值']: 13, C['残值率']: 8,
         C['折旧月数']: 9, C['建账前已提']: 12, C['处置日期']: 11, C['备注']: 18, S['月折旧']: 11, S['累计折旧']: 13,
         S['净值']: 13, S['校验']: 40}
    fm = {C['取得日期']: DATE, C['处置日期']: DATE, C['原值']: MONEY, C['建账前已提']: MONEY, C['残值率']: '0%', C['折旧月数']: '0',
          C['编号']: '@', S['月折旧']: MONEY, S['累计折旧']: MONEY, S['净值']: MONEY}
    al = {C[k]: AL for k in ('名称', '类别', '备注')}
    al.update({C['原值']: AR, C['建账前已提']: AR})
    rows = ctx.get('fas', [])
    _frame(ws, FA_LAST, heads, C_MASTER, tip, W, N_FA, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=C['备注'])
    _kpi(ws, 3, [('A', '原值合计', '=SUMIFS(卡_原值,卡_有效,1)', MONEY),
                 ('D', '本年折旧（到截止日）', '=SUM(卡_月折旧块)', MONEY)])
    _nofx(ws, C['名称'], S['校验'], f'{C["处置日期"]}3', R0 + N_FA + 1000, f'{S["校验"]}3')
    _write_tuples(ws, rows, C, R0, '备注')
    ok = 'INDEX(卡_有效,{n})<>1'
    acc = 'ROUND(INDEX(卡_建账前已提,{n})+SUM(INDEX(卡_月折旧块,{n},0)),2)'
    _show(ws, N_FA, {
        S['月折旧']: f'IF({ok},"",INDEX(卡_月折旧,{{n}}))',
        S['累计折旧']: f'IF({ok},"",{acc})',
        S['净值']: f'IF({ok},"",ROUND(INDEX(卡_原值,{{n}})-{acc},2))',
        S['校验']: _chk1('卡'),
    }, {S['月折旧']: MONEY, S['累计折旧']: MONEY, S['净值']: MONEY},
        {S['月折旧']: AR, S['累计折旧']: AR, S['净值']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_FA - 1)
    rng = rng_of(N_FA)
    dv_fixed(ws, rng(C['类别']), ['机器设备', '运输设备', '电子设备', '房屋建筑', '办公家具'], stop=False,
             error='不是常用类别（可以直接写，点「是」）')
    dv_fixed(ws, rng(C['使用部门']), FA_DEPTS, prompt='生产→制造费用-折旧；管理→管理费用-折旧；销售→销售费用')
    for k in ('取得日期', '处置日期'):
        dv_day(ws, rng(C[k]))
    dv_dec(ws, rng(C['原值']), '原值', prompt='不含税价（专票的税额已经抵了进项）')
    dv_dec(ws, rng(C['残值率']), '残值率', '0', '1', prompt='空＝5%')
    dv_whole(ws, rng(C['折旧月数']), 1, 600, '折旧月数', prompt='机器设备一般 120，运输设备 48，电子设备 36')
    dv_dec(ws, rng(C['建账前已提']), '建账前已提', prompt='建账日以前已经提了多少折旧；空着＝按取得次月到建账日的月数自动算')
    _tail(ws, FA_LAST, '卡片末行', N_FA)


# ───────────────────────── 固定支出 ─────────────────────────
def build_fix(ws, ctx):
    C, S = FIX_COLS, FIX_SHOW
    req = ('类别', '每月金额')
    sub = {'类别': '收支类别', '开始月份': '空＝1月', '结束月份': '空＝12月'}
    hd = {'类别': '收支类别'}
    heads = [(C[k], _hd(hd.get(k, k), tuple(hd.get(x, x) for x in req), sub.get(k, '') if k != '类别' else '')) for k in C] + [
        (S['本月已付'], '本月已付\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 每个月固定要付的（厂房租金、宿舍租金、社保、贷款利息……）一条一行。资金台帐付这笔钱时，编号填这里的编号（如 GD01）；'
           '不填也行，收支类别和往来单位对得上就自动算已付。开始、结束月份填 1～12（或 2026-8）。首页提醒本月还没付的。灰底是示例。')
    W = {C['编号']: 8, C['项目']: 14, C['类别']: 14, C['收款单位']: 13, C['每月金额']: 12, C['每月几号']: 9, C['付款账户']: 11,
         C['开始月份']: 10, C['结束月份']: 10, C['备注']: 20, S['本月已付']: 12, S['校验']: 44}
    fm = {C['每月金额']: MONEY, C['开始月份']: MON_FMT, C['结束月份']: MON_FMT, C['编号']: '@', C['每月几号']: '0',
          S['本月已付']: MONEY}
    al = {C[k]: AL for k in ('项目', '收款单位', '备注')}
    al[C['每月金额']] = AR
    rows = ctx.get('fixed', [])
    _frame(ws, FIX_LAST, heads, C_MASTER, tip, W, N_FIX, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=C['备注'])
    _kpi(ws, 3, [('A', '本月固定支出', '=SUMIFS(固_每月金额,固_本月生效,1)', MONEY),
                 ('D', '本月已付', '=SUMIFS(固_本月已付,固_本月生效,1)', MONEY)])
    put(ws, 'G3', '="本月＝截止日所在月 "&TEXT(P_截止,"yyyy-mm")', F_NOTE, align=AL, border=False)
    _nofx(ws, C['项目'], S['校验'], f'{C["备注"]}3', R0 + N_FIX + 1000, f'{S["校验"]}3')
    _write_tuples(ws, rows, C, R0, '备注')
    _show(ws, N_FIX, {
        S['本月已付']: 'IF(INDEX(固_本月生效,{n})<>1,"",INDEX(固_本月已付,{n}))',
        S['校验']: _chk1('固'),
    }, {S['本月已付']: MONEY}, {S['校验']: AL, S['本月已付']: AR})
    _check_cf(ws, S['校验'], R0, R0 + N_FIX - 1)
    rng = rng_of(N_FIX)
    dv_names(ws, rng(C['类别']), '收支类别列表', '跟资金台帐付这笔钱时选的收支类别一样')
    dv_names(ws, rng(C['收款单位']), '往来单位列表', '可空；填了的话资金台帐的往来单位也要一样才算已付', stop=False)
    dv_dec(ws, rng(C['每月金额']), '每月金额')
    dv_whole(ws, rng(C['每月几号']), 1, 31, '每月几号')
    dv_names(ws, rng(C['付款账户']), '账户列表', '一般从哪个账户付')
    for k in ('开始月份', '结束月份'):
        dv_tip(ws, rng(C[k]), '填 1～12（本年度几月），或 2026-8 这样；空着＝全年')
    dv_len(ws, rng(C['编号']), 12)
    _tail(ws, FIX_LAST, '固定末行', N_FIX)


# ───────────────────────── 会计科目 ─────────────────────────
def build_coa(ws, ctx):
    C, S = COA_COLS, COA_SHOW
    req = ('编码', '名称')
    sub = {'编码': '不要改', '方向': '借/贷', '成本项目': '制造费用明细用'}
    heads = [(C[k], _hd(k, req, sub.get(k, ''))) for k in C] + [(S['全称'], '全称\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('★ 会计用。编码不要改（自动凭证按编码找科目）；可以加二级明细（6 位，前 4 位是一级的编码）、改名称。类别、方向照一级的填。'
           '成本项目：制造费用明细在【成本费用分析】里归到哪一项（水、电、供热、人工、折旧、建厂费用摊销…）。现金、银行、存货、生产成本这些科目由自动凭证记。')
    W = {C['编码']: 9, C['名称']: 18, C['类别']: 7, C['方向']: 6, C['成本项目']: 13, C['说明']: 34, S['全称']: 26, S['校验']: 40}
    fm = {C['编码']: '@'}
    al = {C[k]: AL for k in ('编码', '名称', '成本项目', '说明')}
    rows = ctx.get('coa', [])
    _frame(ws, COA_LAST, heads, C_MASTER, tip, W, N_COA, fm, al, f'C{R0}', auto=tuple(S.values()), title_col=COA_LAST)
    _kpi(ws, 3, [('A', '科目', '=COUNTIF(科目_有效,1)&" 个"', '@')])
    for i, row in enumerate(rows):
        r = R0 + i
        for k, v in zip(C, row):
            v = _v(v)
            if v is not None:
                ws[f'{C[k]}{r}'] = str(v) if k == '编码' else v
    a = f'TRIM({own(C["编码"])}&"")'
    b = f'TRIM({own(C["名称"])}&"")'
    cl = f'TRIM({own(C["类别"])}&"")'
    d = f'TRIM({own(C["方向"])}&"")'
    codes = f'${C["编码"]}${R0}:${C["编码"]}${R0 + N_COA - 1}'
    chk = (f'IF(AND({a}="",{b}=""),"",IF({a}="","✗ 没填编码",IF(ISERROR(--{a}),"✗ 编码只能是数字",'
           f'IF(AND(LEN({a})<>4,LEN({a})<>6),"✗ 编码一级 4 位、二级 6 位",IF(COUNTIF({codes},{a})>1,"✗ 编码重复了",'
           f'IF(AND(LEN({a})=6,COUNTIF({codes},LEFT({a},4))=0),"✗ 上级科目 "&LEFT({a},4)&" 不在表里",'
           f'IF({b}="","✗ 没填名称",IF(NOT(OR({cl}="",' + ','.join(f'{cl}="{x}"' for x in COA_CLASSES) + ')),'
           f'"✗ 类别要选资产、负债、权益、成本、损益",IF(NOT(OR({d}="",{d}="借",{d}="贷")),"⚠ 方向填借或贷（空＝按类别）","✓")))))))))')
    _show(ws, N_COA, {
        S['全称']: f'IF({a}="","",INDEX(科目_全称,{{n}})&"")',
        S['校验']: chk,
    }, {}, {S['全称']: AL, S['校验']: AL})
    _check_cf(ws, S['校验'], R0, R0 + N_COA - 1)
    rng = rng_of(N_COA)
    dv_fixed(ws, rng(C['类别']), COA_CLASSES)
    dv_fixed(ws, rng(C['方向']), ['借', '贷'], prompt='空着：资产、成本、损益＝借；负债、权益＝贷')
    dv_tip(ws, rng(C['编码']), '已有科目的编码不要改；加明细：6 位，前 4 位是一级编码（如 410117）')
    home_link(ws, f'{CL(CI(COA_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')


BUILDERS = {SH_CASH: build_cash, SH_PUR: build_pur, SH_SALE: build_sale, SH_PROD: build_prod, SH_WAGE: build_wage,
            SH_INV: build_inv, SH_CNT: build_cnt, SH_MJ: build_mj,
            SH_BASE: build_base, SH_UNIT: build_unit, SH_ITEM: build_item, SH_FA: build_fa, SH_FIX: build_fix, SH_COA: build_coa}


def build(wb, ctx):
    for name, fn in BUILDERS.items():
        fn(wb[name], ctx)
