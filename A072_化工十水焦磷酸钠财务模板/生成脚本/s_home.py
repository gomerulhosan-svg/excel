# -*- coding: utf-8 -*-
"""A072 化工 · 【首页】：截止日期（全书默认，黄格 layout.HOME_END）、① 钱、② 欠、③ 产销存、④ 成本利润、⑤ 提醒、⑥ 各表入口（会计用的单放）、
   启用三步、常见情况怎么记。

   取数（只用定义名称和本表隐藏格；不引用别的查看表的格子——查看表的黄格改了不影响首页）：
   ① 钱（截至截止日，按日期）：可用资金＝现金/银行/微信支付宝账户（重名的只算第一个，跟资金日报一样）期初＋SUMIFS(资_净额)；承兑＝票据账户；
      今天/本月实际收支＝现金账户、有效、归类不是内部转账（跟资金日报「全公司实际收支」一样）。各账户余额（现金账户、承兑）列在下面。
   ② 欠（截至截止日）：客户欠＝Σ位_应收>0；超期＝位_没回款天数>P_回款天数（往来单位没有账期列，跟【客户往来】一样用参数「回款天数」）；
      欠供应商（不含物流）/欠物流＝Σ位_应付>0 按 位_类型 分；
      借款、垫付、投资：每个出借人 l（_表 出借人清单）取 息_借本金/息_投本金/息_借到截止应计/息_投到截止应计；
        其中「垫付」部分单独算：⑦ 性质＝垫付＋资金台帐 垫付账户的行（资_账户类型＝5）＋还垫付的行（资_原对方号＝224101、不在垫付账户），
        本金＝Σ本金变动，应计利息＝ROUND(借款日利率×((截止＋1)×本金－Σ本金×日期),2)（跟 _息 一样的算法）；
        欠合伙人和个人借款＝出借人类别≠银行 的（借本金－垫付本金、借应计－垫付利息）；银行借款＝类别＝银行 的；合伙人投资款单列（本钱不算欠款，「利息」＝应付利润）。
        四行加起来＝【借款明细及利息】上半的本金、到截止日应计未付合计。
      本月预计要交增值税＝本月销项－本月进项－上月留抵（222102 贷－借、222101 借－贷，按 _矩 整月；留抵＝从建账期初 222101＋222102 借方余额起逐月 MAX(0,…) 递推）；
        负数显示「留抵」。
   ③ 产销存（截至截止日）：截止月月初量（引擎 期初量）＋本月到截止日的进/回收/完工/外购－投料/销售（按隐藏层 日期≤截止）＋本月盘点差（引擎 盘量）；
      主原料＝物料档案类别＝原料的，有库存或近 30 天有投料的排前面（再按物料档案顺序），最多 6 个；
      够用几天＝INT(库存÷日均投料)，日均＝截止日往前 30 天（建账不满 30 天按建账以来天数）投料÷天数；没投料「—」。
   ④ 成本利润（按记账：整月）：上个完整月＝截止月－1（早于建账月显示「—」）、本月＝截止月（还没过完）、本年＝建账月～截止月合计；
      每吨制造成本＝Σ本_完工÷Σ本_产量（所有产品）；每吨完全成本＝每吨制造成本＋（5601＋5602＋5603＋[销货运费进主营成本时]540102）÷销量
      （＝5601 扣销货运费＋5602＋5603＋销货运费）；均价、每吨毛利、产量、销量、净利润、分后利润取 利润表_*（第 m 个＝m 月）；本年的每吨数用合计重算。
   ⑤ 提醒：数用 s_check 给首页的名称（跟【数据校验】一个数）：校验_错数/校验_警数、校验_待补票笔数/金额、校验_待定价进货笔数/金额、
      校验_负库存（任一月末量<0 的种数）、校验_工资没登（月数）、校验_固定支出没付条数/金额（过了日子还没付清）；
      待定价发货、示例行、公司名称本表自己算。每条带超链接。
   ⑥ 导航按 layout.GROUPS（首页除外）；会计组（凭证汇总、记账凭证、科目余额表、手工凭证、年末结转）单独放「会计用」一栏。

   隐藏（AA 列以右，F_HELP 小灰字）：AA/AB 标量；AD..AO 每月一行（第 j 月在第 j+1 行）；AQ..AY 账户；BA..BN 出借人；BP/BQ ⑦；
   BS..CH 原料；CJ..CQ 产品。表保护由 make 按 VIEW_GROUPS 加（黄格不锁）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from layout import *
from common import (YH, F_TITLE, F_TIP, F_HDR, F_TXT, F_TXTB, F_AUTOB, F_NOTE, F_HELP, F_KPI_L, F_KPI_V, F_SEC, F_SEL,
                    fill, cf_bg, put, title, header, section, widths, hide, link, dv_date, print_setup,
                    FILL_SEL, FILL_TOT, BD, AC, ACW, AL, ALW, AR, MONEY, DATE)

LINKF = Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single')
LBL = fill('FFD9E1F2')
WHITE = fill('FFFFFFFF')
GREY = fill('FFF2F2F2')
F_RED = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_OK = Font(name=YH, sz=10, color='FF548235')
QTY = '#,##0.000;[Red]-#,##0.000;"-"'
DAYS = '0" 天"'
KPI_COLS = (('B', 'C', 'D'), ('F', 'G', 'H'), ('J', 'K', 'L'))
LAST = 'L'
SHOW_RM = 6                                   # 主原料最多列几个

# ── 隐藏列 ──
SCL, SCV = 'AA', 'AB'
MO = dict(j='AD', 产量='AE', 销量='AF', 完工='AG', 费运='AH', 收入='AI', 成本='AJ', 净利='AK', 分后='AL', 销项='AM', 进项='AN',
          留抵='AO')
AC_ = dict(a='AQ', 名称='AR', 显示='AS', 类型='AT', 人='AU', 余额='AV', 键='AW', 槽='AX', 垫人='AY')
LD = dict(l='BA', 单位号='BB', 名称='BC', 类别='BD', 日率='BE', 借本金='BF', 投本金='BG', 借应计='BH', 投应计='BI', 垫本金='BJ',
          垫PD='BK', 垫息='BL', 银到期='BM', 他到期='BN')
LO7 = dict(n='BP', 金额日='BQ')
RM = dict(k='BS', 名称='BT', 类别号='BU', 槽='BV', 期初='BW', 进='BX', 回='BY', 投='BZ', 销='CA', 盘='CB', 库存='CC', 近30='CD',
          日均='CE', 天='CF', 键='CG', 显示='CH')
PD = dict(p='CJ', 名称='CK', 期初='CL', 完='CM', 购='CN', 销='CO', 盘='CP', 库存='CQ')

# 隐藏标量
SC_KEYS = ['m', 'm1', 'ok1', '天数', '整月',
           '可用', '承兑', '今收', '今支', '月收', '月支',
           '客欠', '超家', '超额', '预收', '欠供', '欠流', '预付',
           '留抵初', '销项', '进项', '上留', '增值税',
           '垫本', '垫息', '个本', '个息', '银本', '银息', '投本', '投息', '银到期', '个到期', '垫人',
           '产量', '销量', '成品', '料显示数',
           'S产', 'S销', 'S完', 'S费', 'S收', 'S成', 'S净', 'S分',
           '待销数', '待销吨', '示例']
SC = {k: 2 + i for i, k in enumerate(SC_KEYS)}

# ── 版面行号（模型、核对脚本也用） ──
R1 = 7                                        # ① 标题行；R1+1、R1+2 卡片
R_ACC0 = R1 + 3                               # 各账户余额（3 个一行）
R_ACC1 = R_ACC0 + (N_ACC + 2) // 3 - 1
R2 = R_ACC1 + 3                               # ② 标题行；R2+1、R2+2 卡片
R_VAT = R2 + 3
R_LH = R_VAT + 1                              # 借款表头；R_LH+1..+4 四行
R3 = R_LH + 6                                 # ③ 标题行；R3+1 卡片
R_MH = R3 + 2                                 # 主原料表头；R_MH+1..+SHOW_RM
R4 = R_MH + SHOW_RM + 2                       # ④ 标题行
R_TH = R4 + 1                                 # ④ 表头；R_TH+1..+8
N4 = 8
R5 = R_TH + N4 + 3                            # ⑤ 标题行；R5+1.. 提醒


def sc(k):
    return f'${SCV}${SC[k]}'


def rng(col, n):
    return f'${col}$2:${col}${n + 1}'


def hset(ws, coord, f, fmt=None):
    ws[coord] = f if isinstance(f, (int, float)) else '=' + f
    ws[coord].font = F_HELP
    if fmt:
        ws[coord].number_format = fmt


def tx(x, fmt='#,##0.00'):
    return f'TEXT({x},"{fmt}")'


# ── 示例行：各表「备注/说明」列以「示例」开头的行（data.py 的示例都这么写） ──
EX_COLS = [(SH_CASH, CASH_COLS['备注']), (SH_PUR, PUR_COLS['备注']), (SH_SALE, SALE_COLS['备注']), (SH_PROD, PROD_COLS['备注']),
           (SH_WAGE, WAGE_COLS['备注']), (SH_INV, INV_COLS['备注']), (SH_CNT, CNT_COLS['备注']), (SH_MJ, MJ_COLS['备注']),
           (SH_UNIT, UNIT_COLS['备注']), (SH_FA, FA_COLS['备注']), (SH_FIX, FIX_COLS['备注']),
           (SH_BASE, PER_COLS['备注']), (SH_BASE, LO_COLS['说明']), (SH_BASE, OP_COLS['说明'])]

# ── 导航 ──
DESC = {
    SH_CASH: '所有账户的每一笔钱（老板垫付、备用金也记这里），有即时余额', SH_PUR: '进货一车一行：货款＋运费；没定价的先空着单价',
    SH_SALE: '发货一车一行：客户、吨位、单价，我方出的运费', SH_PROD: '一天一个产品一行：产量、废料、各原料投料',
    SH_WAGE: '月底一人一行：应发、扣社保、扣借支', SH_INV: '收到、开出的发票（补票填垫付编号）', SH_CNT: '月底实盘数量（差异自动调账）',
    SH_DAY: '选一天：各账户收支余额、当天每一笔', SH_MON: '选一个月：各账户、各类收支、全年各月', SH_FEE: '按收支类别 × 1～12 月，点类别看每一笔',
    SH_REIM: '谁报了什么、每人每月', SH_ADV: '老板垫付明细（已报/未报）、补票跟进', SH_LOAN: '每个出借人的本金、利息、到期日',
    SH_STK: '选月份：原料、回收料、产品的进出存', SH_COST: '选产品：每吨成本拆项、单耗、收率、毛利',
    SH_PL: '1～12 月利润表＋每吨均价、成本、毛利', SH_BS: '选月份：资产、负债、权益', SH_CF: '1～12 月现金流量',
    SH_AR: '每个客户欠多少、多久没回款、开票', SH_AP: '欠每个供应商、物流多少、欠票', SH_STMT: '选一家出对账单（A4 打印）',
    SH_INVS: '每月开票收票、没开票的客户、欠票的供应商',
    SH_BASE: '参数、账户、收支类别、报销项目、人员、票种、期初', SH_UNIT: '客户、供应商、物流、出借人（期初应收应付、利率）',
    SH_ITEM: '原料、辅料、产品（期初盘点、参考单价）', SH_FA: '设备卡片（每月自动折旧）', SH_FIX: '每月固定要付的（房租等）',
    SH_COA: '会计科目（编码别改，可以加二级）', SH_CHK: '哪里录错了、账对不对、该跟进什么',
    SH_VLIST: '每月凭证清单（张数、页数）', SH_VPRINT: '记账凭证打印（A4 一张两页）', SH_TB: '选月份：每个科目期初、发生、期末',
    SH_MJ: '会计调账用的手工凭证', SH_YE: '年末数整理好，抄到明年新本的期初',
}
GRP_NAME = {'录入': '平时录（蓝）', '经营': '经营（绿，自动）', '报表': '报表（紫，自动）', '往来': '往来（橙，自动）',
            '基础': '基础（灰，开账填、偶尔改）', '校验': '校验'}

TIP = ('★ 平时只录蓝色的几张：【资金台帐】（每一笔钱，老板垫付、备用金也在这里记）、【进货物流登记】（一车一行：货款＋运费）、【销售台账】（一车一行）、'
       '【生产日报】，月底【工资表】【月末盘点】，收到发票登【发票登记】；凭证、报表、成本、往来、借款利息都是自动的。录入表都可以插行、删行、排序。')

STEPS = [
    '① 删示例：各表灰底、备注写着「示例」的行整行删掉（基础资料⑤人员、⑦⑧，往来单位、固定资产、固定支出里的示例也删）；'
    '收支类别、报销项目、会计科目是预置好的，按自己的情况改名、增删。',
    '② 填期初（建账日前一天晚上的数）：【基础资料】公司名称、建账日（某月 1 号）、② 各账户期初余额（垫付、备用金账户不填）、'
    '⑦ 建账前的借款/投资/老板垫付一笔一行、⑧ 其他期初（押金、没摊完的建厂费用、欠的工资税费…）；【往来单位】客户期初应收、供应商期初应付、'
    '出借人的利率和到期日；【物料档案】期初盘点数量和单价；【固定资产】一台设备一张卡片。填完看【数据校验】期初试算平衡是 √。',
    '③ 每天录：从建账日起每天录蓝色的表，月底登【工资表】和【月末盘点】；每天看一眼首页，有 ✗ 就去【数据校验】看是哪一行。',
]

HOWTO = [
    ('客户打来货款、收到承兑', '【资金台帐】账户选收到钱的账户（承兑选「承兑汇票」账户），类别「收客户货款」，往来单位选客户，填收入；退给客户的记支出'),
    ('承兑到期、贴现', '记两行，类别都选「承兑到期/贴现入账」：承兑汇票账户记支出（票面）、银行记收入（到账的）；贴现少的钱在承兑汇票账户记一行「贴现息」'),
    ('进一车原料', '【进货物流登记】一行：物料、供应商、数量、单价（还没定价的单价、金额都空着＝按物料档案参考单价暂估，定了价回原行补）；'
                '运费填「运费」或「每吨运费」，物流单位空着＝供应商。退货数量填负数；扣款折让数量空着、货款金额填负数'),
    ('付原料货款、付运费', '【资金台帐】类别「原材料材料款」（付运费选「原材料运输费」或「销货运费」），往来单位选供应商或物流单位：冲进货、销售里欠下的'),
    ('发一车货', '【销售台账】一行：客户、产品（空着＝默认产品）、单价、吨位；我方出运费的填物流单位、运费；开了票填发票号。'
              '还没定价的单价空着（出库照算，收入定价后进）；卖废料、卖原料：产品选那种料'),
    ('生产', '【生产日报】一天一个产品一行：产量（袋）、废料（袋）、各原料投了多少吨（列名就是原料名，在物料档案「生产日报第几列」设）'),
    ('老板自己掏钱付的', '【资金台帐】账户选「老板垫付」，类别选真实用途（伙食费、员工报销…），票种空着＝待补票，编号填资金台帐第 3 行提示的下一个垫付编号。'
                  '公司还老板：只在银行记一行，类别「垫付现金已报」，往来单位选老板'),
    ('补到发票', '【发票登记】一行，方向「收到」，垫付编号填那一笔的编号：专票自动抵进项，待补票就消掉了'),
    ('员工报销、备用金', '类别「员工报销」，报销人、报销项目必填。给员工备用金：银行记支出，类别「发备用金」，往来单位选这个人；'
                  '他用备用金买东西，在他的备用金账户记支出'),
    ('借款、还款、付利息', '出借人先在【往来单位】建（类型选出借人，填类别、利率、到期日）。借进来选「借款到账」、还本金选「还出」，往来单位选出借人；'
                  '付利息选「农行利息」「微众贷款利息」这类；利息每月自动计提'),
    ('合伙人投资', '类别「投资款」「退还投资款」；按投资年利率算的「利息」每月自动计进利润分配，付的时候选「付投资款利息」'),
    ('工资', '月底【工资表】一人一行（应发、扣社保、扣借支），自动计提；发工资时【资金台帐】类别「员工工费」，预支选「员工借支」'),
    ('电费、水费、蒸汽（下个月才付）', '付款那一行「所属月份」填用的那个月（填 8 就行），成本算进那个月'),
    ('一次付几个月房租、建厂费用', '环评、认证、装修、开办费这些类别已经设好自动分月摊；一次付一年房租的，新建一个类别（如「厂房租金（年付）」）分摊月数填 12'),
    ('买设备', '类别「机械设备费」，再到【固定资产】登一张卡片（原值填不含税），每月自动折旧'),
    ('自己账户之间倒钱', '记两行「内部转账」：转出账户记支出、转进账户记收入，金额一样'),
    ('缴税', '选「缴税-增值税」「缴税-附加税印花税」「缴税-所得税」「缴税-代扣个税」，一张缴款书几种税就拆几行'),
    ('月底', '登【工资表】、【月末盘点】（填实盘数量，差异自动调账）；看【进销存】【成本费用分析】【利润表】；有 ✗ 去【数据校验】'),
    ('调账', '会计在【手工凭证】做（同一个组号借贷要平）；平时不用'),
]


def notes():
    return [
        '① 这本表管一个自然年（建账日那一年）。明年用【年末结转】把年末数抄到新本的期初。',
        '② 客户、供应商、物料、账户、收支类别都从下拉选；新的先到【往来单位】【物料档案】【基础资料】加一行（全称、别名也认）。不要改表名（链接会失效）。',
        '③ 录入表、基础资料都可以插行、删行、排序（整行一起动）；绿、紫、橙色的表和首页是自动的，有保护（没有密码），只有黄格子能改。',
        f'④ 资金台帐最多 {N_CASH} 笔、进货和销售各 {N_PUR}/{N_SALE} 车、生产日报 {N_PROD} 行、工资表 {N_WAGE} 行、发票 {N_INV} 张、'
        f'手工凭证 {N_MJ} 行；超了【数据校验】会提示。',
        '⑤ 会计用的几张（凭证汇总、记账凭证、科目余额表、手工凭证、年末结转）给会计看、打印，老板平时不用管。',
    ]


# ═══════════════════════════════ 隐藏计算 ═══════════════════════════════
def build_hidden(ws):
    for k, r in SC.items():
        ws[f'{SCL}{r}'] = k
        ws[f'{SCL}{r}'].font = F_HELP
    S = lambda k, f, fmt=None: hset(ws, f'{SCV}{SC[k]}', f, fmt)
    m = sc('m')
    # 标签行
    for d in (MO, AC_, LD, LO7, RM, PD):
        for k, c in d.items():
            ws[f'{c}1'] = k
            ws[f'{c}1'].font = F_HELP

    # ── 每月（第 j 月在第 j+1 行） ──
    k1, k2 = kh('222101'), kh('222102')
    hb = lambda code, j: f'(INDEX(矩_汇借,{kh(code)},{j})-INDEX(矩_汇贷,{kh(code)},{j}))'
    for j in MONTHS:
        r = j + 1
        c = lambda k: f'{MO[k]}{r}'
        hset(ws, c('j'), j)
        hset(ws, c('产量'), f'ROUND(N(INDEX(利润表_产量,{j})),3)')
        hset(ws, c('销量'), f'ROUND(N(INDEX(利润表_销量,{j})),3)')
        hset(ws, c('完工'), f'ROUND(SUM(INDEX(本_完工,0,{j})),2)')
        hset(ws, c('费运'), f'ROUND({hb("5601", j)}+{hb("5602", j)}+{hb("5603", j)}+IF(P_运费进成本=1,{hb("540102", j)},0),2)')
        hset(ws, c('收入'), f'ROUND(-{hb("5001", j)},2)')
        hset(ws, c('成本'), f'ROUND({hb("5401", j)},2)')
        hset(ws, c('净利'), f'ROUND(N(INDEX(利润表_净利润,{j})),2)')
        hset(ws, c('分后'), f'ROUND(N(INDEX(利润表_分后利润,{j})),2)')
        hset(ws, c('销项'), f'ROUND(INDEX(矩_汇贷,{k2},{j})-INDEX(矩_汇借,{k2},{j}),2)')
        hset(ws, c('进项'), f'ROUND(INDEX(矩_汇借,{k1},{j})-INDEX(矩_汇贷,{k1},{j}),2)')
        prev = sc('留抵初') if j == 1 else f'{MO["留抵"]}{r - 1}'
        hset(ws, c('留抵'), f'ROUND(MAX(0,{prev}+{c("进项")}-{c("销项")}),2)')
    mo = lambda k: f'${MO[k]}$2:${MO[k]}$13'
    at = lambda k, j: f'INDEX({mo(k)},{j})'
    upto = lambda k: f'SUM(${MO[k]}$2:INDEX({mo(k)},{m}))'

    # ── 账户（第 a 个在第 a+1 行） ──
    A = AC_
    for a in range(1, N_ACC + 1):
        r = a + 1
        c = lambda k: f'{A[k]}{r}'
        hset(ws, c('a'), a)
        hset(ws, c('名称'), f'INDEX(账户_名称,{a})&""')
        hset(ws, c('显示'), f'IF({c("名称")}="",0,IF(INDEX(账户_类型号,{a})=0,0,IF(MATCH({esc(c("名称"))},账户_名称,0)={a},1,0)))')
        hset(ws, c('类型'), f'IF({c("显示")}=1,INDEX(账户_类型号,{a}),0)')
        pn = f'TRIM(INDEX(账户_人员,{a})&"")'
        hset(ws, c('人'), f'IF({c("类型")}=5,IF({pn}="",{c("名称")},{pn}),"")')
        hset(ws, c('余额'), (f'IF(OR({c("类型")}<1,{c("类型")}>4),0,ROUND(INDEX(账户_期初,{a})'
                             f'+SUMIFS(资_净额,资_账户号,{a},资_日期,"<="&P_截止),2))'), MONEY)
        hset(ws, c('键'), f'IF(AND({c("类型")}>=1,{c("类型")}<=4),IF({c("类型")}<=3,1,2)*100+{a},"")')
        hset(ws, c('槽'), f'IF({a}>COUNT({rng(A["键"], N_ACC)}),0,MOD(SMALL({rng(A["键"], N_ACC)},{a}),100))')
        hset(ws, c('垫人'), f'IF({c("人")}="","",IF(COUNTIF(${A["人"]}$2:{c("人")},{esc(c("人"))})=1,{c("人")},""))')
    acc = lambda k: rng(A[k], N_ACC)

    # ── ⑦（第 n 条在第 n+1 行）：金额×日期 ──
    for n in range(1, N_LO + 1):
        r = n + 1
        hset(ws, f'{LO7["n"]}{r}', n)
        hset(ws, f'{LO7["金额日"]}{r}', f'INDEX(借期_金额,${LO7["n"]}{r})*INDEX(借期_日期,${LO7["n"]}{r})')
    lo_ad = rng(LO7['金额日'], N_LO)

    # ── 出借人（第 l 个在第 l+1 行） ──
    KD = kh('224101')
    for l in range(1, N_LEND + 1):
        r = l + 1
        c = lambda k: f'{LD[k]}{r}'
        L = f'${LD["l"]}{r}'
        u = c('单位号')
        hset(ws, c('l'), l)
        hset(ws, u, f'N(INDEX(出借_单位号,{L}))')
        hset(ws, c('名称'), f'INDEX(出借_名称,{L})&""')
        hset(ws, c('类别'), f'IF({u}=0,0,INDEX(位_出借人类别,{u}))')
        hset(ws, c('日率'), f'IF({u}=0,0,INDEX(位_借款利率,{u})/INDEX(位_计息天数,{u}))')
        for k, nm in (('借本金', '息_借本金'), ('投本金', '息_投本金'), ('借应计', '息_借到截止应计'), ('投应计', '息_投到截止应计')):
            hset(ws, c(k), f'N(INDEX({nm},{L}))', MONEY)
        lo = f'借期_出借人号,{L},借期_性质,3,借期_有效,1'
        z5 = f'资_出借人号,{L},资_账户类型,5,资_日期,"<="&P_截止'
        zb = f'资_出借人号,{L},资_原对方号,{KD},资_账户类型,"<>5",资_日期,"<="&P_截止'
        hset(ws, c('垫本金'), (f'IF({u}=0,0,ROUND(SUMIFS(借期_金额,{lo})+SUMIFS(资_本金变动,{z5})+SUMIFS(资_本金变动,{zb}),2))'), MONEY)
        hset(ws, c('垫PD'), f'IF({u}=0,0,SUMIFS({lo_ad},{lo})+SUMIFS(资_本金乘日,{z5})+SUMIFS(资_本金乘日,{zb}))')
        hset(ws, c('垫息'), f'IF({c("日率")}=0,0,ROUND({c("日率")}*((P_截止+1)*{c("垫本金")}-{c("垫PD")}),2))', MONEY)
        due = f'INDEX(位_到期日,MAX(1,{u}))'
        pos = f'{u}>0,{c("借本金")}-{c("垫本金")}>0.005,{due}>0'
        hset(ws, c('银到期'), f'IF(AND({pos},{c("类别")}=1),{due},"")', DATE)
        hset(ws, c('他到期'), f'IF(AND({pos},{c("类别")}<>1),{due},"")', DATE)
    ld = lambda k: rng(LD[k], N_LEND)

    # ── 原料（第 k 个在第 k+1 行）：截至截止日的库存、近 30 天投料 ──
    R_ = RM
    d0 = 'P_截止'
    slot_now = lambda j: f'SUMIFS(产_槽{j},产_有效,1,产_月,{m},产_日期,"<="&{d0})'
    slot_30 = lambda j: f'SUMIFS(产_槽{j},产_有效,1,产_日期,">="&({d0}-29),产_日期,"<="&{d0})'
    for k in range(1, N_RM + 1):
        r = k + 1
        c = lambda f: f'{R_[f]}{r}'
        K = f'${R_["k"]}{r}'
        key = f'{K}*100+{m}'
        hset(ws, c('k'), k)
        hset(ws, c('名称'), f'INDEX(料档_名称,{K})&""')
        hset(ws, c('类别号'), f'N(INDEX(料档_类别号,{K}))')
        hset(ws, c('槽'), f'IFERROR(MATCH({K},槽_料号,0),0)')
        hset(ws, c('期初'), f'N(INDEX(料_期初量,{K},{m}))', QTY)
        hset(ws, c('进'), f'SUMIFS(进_数量,进_料月键,{key},进_日期,"<="&{d0})', QTY)
        hset(ws, c('回'), f'SUMIFS(产_废料吨,产_回月键,{key},产_日期,"<="&{d0})', QTY)
        j = c('槽')
        hset(ws, c('投'), f'IF({j}=0,0,CHOOSE({j},{",".join(slot_now(i) for i in range(1, N_SLOT + 1))}))', QTY)
        hset(ws, c('销'), f'SUMIFS(销_吨位,销_出库键,{key},销_出库,2,销_日期,"<="&{d0})', QTY)
        hset(ws, c('盘'), f'N(INDEX(料_盘量,{K},{m}))', QTY)
        hset(ws, c('库存'), f'ROUND({c("期初")}+{c("进")}+{c("回")}-{c("投")}-{c("销")}+{c("盘")},3)', QTY)
        hset(ws, c('近30'), f'IF({j}=0,0,ROUND(CHOOSE({j},{",".join(slot_30(i) for i in range(1, N_SLOT + 1))}),3))', QTY)
        hset(ws, c('日均'), f'ROUND({c("近30")}/{sc("天数")},3)', QTY)
        hset(ws, c('天'), f'IF({c("日均")}<=0,"—",MAX(0,INT(ROUND({c("库存")}/{c("日均")},6))))')
        act = f'OR(ABS({c("库存")})>0.0005,{c("近30")}>0)'
        hset(ws, c('键'), f'IF(AND({c("类别号")}=1,{c("名称")}<>""),IF({act},0,1)*1000+{K},"")')
    for i in range(1, N_RM + 1):
        r = i + 1
        keys = rng(R_['键'], N_RM)
        hset(ws, f'{R_["显示"]}{r}', f'IF({i}>COUNT({keys}),0,MOD(SMALL({keys},{i}),1000))')

    # ── 产品（第 p 个在第 p+1 行） ──
    for p in range(1, N_PRD + 1):
        r = p + 1
        c = lambda f: f'{PD[f]}{r}'
        Pp = f'${PD["p"]}{r}'
        key = f'{Pp}*100+{m}'
        hset(ws, c('p'), p)
        hset(ws, c('名称'), f'INDEX(品档_名称,{Pp})&""')
        hset(ws, c('期初'), f'N(INDEX(品_期初量,{Pp},{m}))', QTY)
        hset(ws, c('完'), f'SUMIFS(产_成品吨,产_品月键,{key},产_日期,"<="&{d0})', QTY)
        hset(ws, c('购'), f'SUMIFS(进_数量,进_品月键,{key},进_日期,"<="&{d0})', QTY)
        hset(ws, c('销'), f'SUMIFS(销_吨位,销_出库键,{key},销_出库,1,销_日期,"<="&{d0})', QTY)
        hset(ws, c('盘'), f'N(INDEX(品_盘量,{Pp},{m}))', QTY)
        hset(ws, c('库存'), f'ROUND({c("期初")}+{c("完")}+{c("购")}-{c("销")}+{c("盘")},3)', QTY)

    # ── 标量 ──
    S('m', 'P_截止月')
    S('m1', 'P_截止月-1')
    S('ok1', f'IF(AND({sc("m1")}>=1,{sc("m1")}>=P_建账月),1,0)')
    S('天数', 'MIN(30,P_截止-P_建账日+1)')
    S('整月', 'IF(P_截止=DATE(P_年度,P_截止月+1,0),1,0)')
    S('可用', f'ROUND(SUMIFS({acc("余额")},{acc("类型")},">=1",{acc("类型")},"<=3"),2)', MONEY)
    S('承兑', f'ROUND(SUMIFS({acc("余额")},{acc("类型")},4),2)', MONEY)
    cash = '资_有效,1,资_现金,1,资_归类,"<>内部转账"'
    S('今收', f'ROUND(SUMIFS(资_收入,资_日期,P_截止,{cash}),2)', MONEY)
    S('今支', f'ROUND(SUMIFS(资_支出,资_日期,P_截止,{cash}),2)', MONEY)
    mon_ = '资_日期,">="&P_月初,资_日期,"<="&P_截止'
    S('月收', f'ROUND(SUMIFS(资_收入,{mon_},{cash}),2)', MONEY)
    S('月支', f'ROUND(SUMIFS(资_支出,{mon_},{cash}),2)', MONEY)
    S('客欠', 'ROUND(SUMIFS(位_应收,位_有效,1,位_应收,">0"),2)', MONEY)
    od = '位_有效,1,位_没回款天数,">"&P_回款天数'
    S('超家', f'COUNTIFS({od})')
    S('超额', f'ROUND(SUMIFS(位_应收,{od}),2)', MONEY)
    S('预收', 'ROUND(-SUMIFS(位_应收,位_有效,1,位_应收,"<0"),2)', MONEY)
    S('欠供', 'ROUND(SUMIFS(位_应付,位_有效,1,位_应付,">0",位_类型,"<>物流"),2)', MONEY)
    S('欠流', 'ROUND(SUMIFS(位_应付,位_有效,1,位_应付,">0",位_类型,"物流"),2)', MONEY)
    S('预付', 'ROUND(-SUMIFS(位_应付,位_有效,1,位_应付,"<0"),2)', MONEY)
    S('留抵初', f'MAX(0,ROUND(INDEX(矩_汇期初,{k1})+INDEX(矩_汇期初,{k2}),2))', MONEY)
    S('销项', at('销项', m), MONEY)
    S('进项', at('进项', m), MONEY)
    S('上留', f'IF({m}=1,{sc("留抵初")},{at("留抵", f"{m}-1")})', MONEY)
    S('增值税', f'ROUND({sc("销项")}-{sc("进项")}-{sc("上留")},2)', MONEY)
    S('垫本', f'ROUND(SUM({ld("垫本金")}),2)', MONEY)
    S('垫息', f'ROUND(SUM({ld("垫息")}),2)', MONEY)
    for key, crit in (('个', '"<>1"'), ('银', '1')):
        S(f'{key}本', f'ROUND(SUMIFS({ld("借本金")},{ld("类别")},{crit})-SUMIFS({ld("垫本金")},{ld("类别")},{crit}),2)', MONEY)
        S(f'{key}息', f'ROUND(SUMIFS({ld("借应计")},{ld("类别")},{crit})-SUMIFS({ld("垫息")},{ld("类别")},{crit}),2)', MONEY)
    S('投本', f'ROUND(SUM({ld("投本金")}),2)', MONEY)
    S('投息', f'ROUND(SUM({ld("投应计")}),2)', MONEY)
    S('银到期', f'MIN({ld("银到期")})', DATE)
    S('个到期', f'MIN({ld("他到期")})', DATE)
    S('垫人', 'MID(' + '&'.join(f'IF({A["垫人"]}{a + 1}="","","、"&{A["垫人"]}{a + 1})' for a in range(1, N_ACC + 1)) + ',2,80)')
    S('产量', f'ROUND(SUM({rng(PD["完"], N_PRD)}),3)', QTY)
    S('销量', f'ROUND(SUM({rng(PD["销"], N_PRD)}),3)', QTY)
    S('成品', f'ROUND(SUM({rng(PD["库存"], N_PRD)}),3)', QTY)
    S('料显示数', f'MIN({SHOW_RM},COUNT({rng(R_["键"], N_RM)}))')
    for k, col in (('S产', '产量'), ('S销', '销量'), ('S完', '完工'), ('S费', '费运'), ('S收', '收入'), ('S成', '成本'),
                   ('S净', '净利'), ('S分', '分后')):
        S(k, f'ROUND({upto(col)},{3 if col in ("产量", "销量") else 2})')
    pend = 'P_截止'
    S('待销数', f'COUNTIFS(销_待定价,1,销_有效,1,销_日期,"<="&{pend})')
    S('待销吨', f'ROUND(SUMIFS(销_吨位,销_待定价,1,销_有效,1,销_日期,"<="&{pend}),3)', QTY)
    S('示例', '+'.join(f"COUNTIF('{sh}'!${col}:${col},\"示例*\")" for sh, col in EX_COLS))


# ═══════════════════════════════ 版面 ═══════════════════════════════
def card(ws, r, i, label, f, fmt=MONEY, big=False, to=None, lab_font=None):
    c0, c1, c2 = KPI_COLS[i]
    lc = put(ws, f'{c0}{r}', label, lab_font or (LINKF if to else F_KPI_L), LBL, align=ALW)
    if to:
        link(lc, to)
    put(ws, f'{c1}{r}', '=' + f, F_KPI_V if big else F_AUTOB, WHITE, fmt, AR)
    put(ws, f'{c2}{r}', None, None, WHITE)
    ws.merge_cells(f'{c1}{r}:{c2}{r}')


def line(ws, r, f, to=None, h=None, font=F_TXT):
    put(ws, f'B{r}', f, font, align=ALW, border=False)
    ws.merge_cells(f'B{r}:{LAST}{r}')
    if to:
        link(ws[f'B{r}'], to)
    if h:
        ws.row_dimensions[r].height = h


def build(wb, ctx=None):
    ws = wb[SH_HOME]
    widths(ws, {'A': 2, 'B': 26, 'C': 15, 'D': 15, 'E': 3, 'F': 24, 'G': 15, 'H': 15, 'I': 3, 'J': 22, 'K': 15, 'L': 15})
    title(ws, '十 水 焦 磷 酸 钠 生 产 销 售 财 务', LAST, C_HOME, TIP)
    ws['A1'].value = '=IF(P_公司名称="","",P_公司名称&" · ")&"十水焦磷酸钠生产销售财务"'
    ws['A2'].value = f'="{TIP}"&IF({sc("示例")}>0,"灰底、备注写着「示例」的行是示例，正式用之前整行删掉。","")'
    build_hidden(ws)

    # ── 截止日期 ──
    put(ws, 'B4', '截止日期', F_KPI_L, LBL, align=AC)
    put(ws, HOME_END, None, F_SEL, FILL_SEL, DATE, AC)
    dv_date(ws, HOME_END)
    ws.row_dimensions[4].height = 24
    last = ('IFERROR("；最后一笔在资金台帐第 "&INDEX(资_录入行,MATCH(P_最后日期,资_日期,0))&" 行",'
            'IFERROR("；最后一笔在进货物流登记第 "&INDEX(进_录入行,MATCH(P_最后日期,进_日期,0))&" 行",'
            'IFERROR("；最后一笔在销售台账第 "&INDEX(销_录入行,MATCH(P_最后日期,销_日期,0))&" 行",'
            'IFERROR("；最后一笔在生产日报第 "&INDEX(产_录入行,MATCH(P_最后日期,产_日期,0))&" 行",""))))')
    put(ws, 'D4', (f'="现在按 "&TEXT(P_截止,"yyyy-mm-dd")&" 算（空着＝最后一笔 "&TEXT(P_最后日期,"yyyy-mm-dd")&"）；建账日 "'
                   f'&TEXT(P_建账日,"yyyy-mm-dd")&{last}'), Font(name=YH, sz=11, bold=True, color='FF1F3864'), align=AL, border=False)
    ws.merge_cells(f'D4:{LAST}4')
    h = HOME_END
    put(ws, 'B5', (f'=IF(AND(TRIM({h}&"")<>"",NOT(ISNUMBER({h}))),"⚠ 截止日期看不懂（要填日期，像 2026/9/30），先按最后一笔的日期算",'
                   f'IF(AND(ISNUMBER({h}),INT({h})<P_建账日),"⚠ 截止日期早于建账日，按建账日 "&TEXT(P_建账日,"yyyy-mm-dd")&" 算",'
                   f'IF(AND(ISNUMBER({h}),INT({h})>P_年末),"⚠ 截止日期晚于 "&P_年度&" 年底，按 "&TEXT(P_年末,"yyyy-mm-dd")&" 算（一本只管一年）",'
                   f'IF(AND(ISNUMBER({h}),P_截止<P_最后日期),"⚠ 截止日期比最新录入（"&TEXT(P_最后日期,"yyyy-mm-dd")&"）早，'
                   f'后面录的都没算进来——是不是忘了清空？",'
                   f'IF(AND(NOT(ISNUMBER({h})),P_截止被挡=1),"⚠ 有 "&P_近月笔数&" 笔日期比别的晚一个多月，多半年份填错了，截止日期没跟过去（看【数据校验】）",'
                   f'"截止日期决定：余额、欠款、库存算到哪天，「今天」「本月」是哪天哪个月。平时空着就行（＝最后一笔的日期）；'
                   f'想看以前某一天的，在黄格里填日期（看完清空）。")))))'), F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'B5:{LAST}5')
    ws.row_dimensions[5].height = 30
    ws.conditional_formatting.add('B5', FormulaRule(formula=['LEFT($B$5,1)="⚠"'], font=F_RED))

    # ── ① 钱 ──
    r = R1
    section(ws, r, 'B', LAST, '① 钱（截至截止日期，按日期算；点蓝字看资金日报、资金月报）', C_HOME)
    card(ws, r + 1, 0, '可用资金\n（现金、银行、微信支付宝）', sc('可用'), big=True, to=SH_DAY)
    card(ws, r + 1, 1, '今天实际收入\n（不含内部转账）', sc('今收'), to=SH_DAY)
    card(ws, r + 1, 2, '本月实际收入\n（1 号到截止日）', sc('月收'), to=SH_MON)
    card(ws, r + 2, 0, '手上承兑汇票\n（不算可用资金）', sc('承兑'), to=SH_DAY)
    card(ws, r + 2, 1, '今天实际支出\n（不含内部转账）', sc('今支'), to=SH_DAY)
    card(ws, r + 2, 2, '本月实际支出\n（1 号到截止日）', sc('月支'), to=SH_MON)
    for rr in (r + 1, r + 2):
        ws.row_dimensions[rr].height = 32
    r0 = R_ACC0
    A = AC_
    for i in range(N_ACC):
        c0, cv, c2 = KPI_COLS[i % 3]
        rr = r0 + i // 3
        a = f'${A["槽"]}${i + 2}'
        an = f'INDEX({rng(A["名称"], N_ACC)},{a})'
        nm = f'{an}&IF(AND(INDEX({rng(A["类型"], N_ACC)},{a})=4,ISERROR(FIND("承兑",{an}))),"（承兑）","")'
        put(ws, f'{c0}{rr}', f'=IF({a}=0,"",{nm})', F_TXTB, align=AL, border=False)
        put(ws, f'{cv}{rr}', f'=IF({a}=0,"",INDEX({rng(A["余额"], N_ACC)},{a}))', F_AUTOB, None, MONEY, AR, border=False)
        ws.merge_cells(f'{cv}{rr}:{c2}{rr}')
    last_acc = R_ACC1
    ws.conditional_formatting.add(f'B{r0}:{LAST}{last_acc}', FormulaRule(formula=[f'AND(ISNUMBER(B{r0}),B{r0}<-0.005)'], font=F_RED))
    for c0, _cv, c2 in KPI_COLS:
        ws.conditional_formatting.add(f'{c0}{r0}:{c2}{last_acc}', FormulaRule(formula=[f'${c0}{r0}<>""'], border=BD))
        ws.conditional_formatting.add(f'{c0}{r0}:{c0}{last_acc}', FormulaRule(formula=[f'${c0}{r0}<>""'], fill=cf_bg('FFF2F2F2')))
    put(ws, f'B{last_acc + 1}', '↑ 各现金账户、承兑汇票账户的余额（截至截止日期）；老板垫付、备用金账户看【资金日报】【个人垫付现金明细】。',
        F_NOTE, align=AL, border=False)
    ws.merge_cells(f'B{last_acc + 1}:{LAST}{last_acc + 1}')

    # ── ② 欠 ──
    r = R2
    section(ws, r, 'B', LAST, '② 欠（截至截止日期；点蓝字看明细）', C_WL)
    card(ws, r + 1, 0, '客户欠我们', sc('客欠'), big=True, to=SH_AR)
    put(ws, f'F{r + 1}', f'="超 "&P_回款天数&" 天没回款的客户"', LINKF, LBL, align=ALW)
    link(ws[f'F{r + 1}'], SH_AR)
    put(ws, f'G{r + 1}', f'={sc("超家")}&" 家 / "&TEXT({sc("超额")},"#,##0.00")', F_AUTOB, WHITE, '@', AR)
    put(ws, f'H{r + 1}', None, None, WHITE)
    ws.merge_cells(f'G{r + 1}:H{r + 1}')
    card(ws, r + 1, 2, '预收客户的（先收了钱）', sc('预收'), to=SH_AR)
    card(ws, r + 2, 0, '欠供应商（不含物流）', sc('欠供'), big=True, to=SH_AP)
    card(ws, r + 2, 1, '欠物流运费', sc('欠流'), to=SH_AP)
    card(ws, r + 2, 2, '预付供应商的', sc('预付'), to=SH_AP)
    vr = R_VAT
    put(ws, f'B{vr}', '本月预计要交增值税', LINKF, LBL, align=ALW)
    link(ws[f'B{vr}'], SH_INVS)
    vat = sc('增值税')
    put(ws, f'C{vr}', f'=IF({vat}>=0,{vat},"留抵 "&TEXT(-{vat},"#,##0.00"))', F_KPI_V, WHITE, MONEY, AR)
    put(ws, f'D{vr}', None, None, WHITE)
    ws.merge_cells(f'C{vr}:D{vr}')
    put(ws, f'F{vr}', (f'="＝"&{sc("m")}&" 月销项 "&TEXT({sc("销项")},"#,##0.00")&" － 进项 "&TEXT({sc("进项")},"#,##0.00")'
                       f'&" － 上月留抵 "&TEXT({sc("上留")},"#,##0.00")&IF({vat}<0,"（进项多，留到下月抵）","")'
                       f'&"。按记账凭证整月算（含补票进项），附加税另算；小规模＝按征收率的应交数"'), F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'F{vr}:{LAST}{vr}')
    for rr in (r + 1, r + 2, vr):
        ws.row_dimensions[rr].height = 30
    lr = R_LH
    header(ws, lr, [('B', '借款、垫付、投资'), ('C', '本金'), ('D', '应计未付利息'), ('F', '合计'), ('G', '说明')], C_WL, height=22)
    put(ws, f'H{lr}', None, None, fill(C_WL))
    ws.merge_cells(f'G{lr}:{LAST}{lr}')
    due_txt = lambda d: (f'IF({d}=0,"",IF({d}<P_截止,"已过期 "&(P_截止-{d})&" 天：","")&"最近 "&TEXT({d},"yyyy-mm-dd")&" 到期"'
                         f'&IF({d}>=P_截止,"（还有 "&({d}-P_截止)&" 天）",""))')
    rows_l = [
        (f'="欠"&IF({sc("垫人")}="","老板",{sc("垫人")})&"垫付"', '垫本', '垫息', SH_ADV,
         f'="垫付账户的支出－公司已还（银行记「垫付现金已报」）＋建账前⑦；利息按出借人借款年利率。下一个垫付编号 "&P_下一垫付号'),
        ('欠合伙人和个人借款', '个本', '个息', SH_LOAN, f'={due_txt(sc("个到期"))}&IF({sc("个到期")}=0,"","；")&"不含上面的垫付"'),
        ('银行借款', '银本', '银息', SH_LOAN,
         f'=IF({sc("银到期")}=0,IF(ABS({sc("银本")})<0.005,"没有银行借款","到期日没填（在【往来单位】填）"),{due_txt(sc("银到期"))})'
         f'&"；利率空着的银行，利息扣了直接进财务费用、不计提"'),
        ('合伙人投资款（本钱，不算欠款）', '投本', '投息', SH_LOAN, '="利息栏＝按投资年利率算、还没付的「利息」（应付利润，从利润里分）"'),
    ]
    for i, (lab, kb, ki, to, note) in enumerate(rows_l):
        rr = lr + 1 + i
        link(put(ws, f'B{rr}', lab, LINKF, LBL, align=ALW), to)
        put(ws, f'C{rr}', f'={sc(kb)}', F_AUTOB, WHITE, MONEY, AR)
        put(ws, f'D{rr}', f'={sc(ki)}', F_AUTOB, WHITE, MONEY, AR)
        put(ws, f'F{rr}', f'=ROUND({sc(kb)}+{sc(ki)},2)', F_TXTB, FILL_TOT, MONEY, AR)
        put(ws, f'G{rr}', note, F_NOTE, align=ALW)
        for cc in 'HIJKL':
            put(ws, f'{cc}{rr}', None)
        ws.merge_cells(f'G{rr}:{LAST}{rr}')
        ws.row_dimensions[rr].height = 30
    ws.conditional_formatting.add(f'G{lr + 1}:G{lr + 3}', FormulaRule(formula=[f'ISNUMBER(SEARCH("已过期",G{lr + 1}))'], font=F_RED))
    ws.conditional_formatting.add(f'C{r + 1}:{LAST}{lr + 4}', FormulaRule(formula=[f'AND(ISNUMBER(C{r + 1}),C{r + 1}<-0.005)'],
                                                                          font=F_RED))

    # ── ③ 产销存 ──
    r = R3
    section(ws, r, 'B', LAST, '③ 产销存（截至截止日期；吨；点蓝字看明细）', C_VIEW)
    card(ws, r + 1, 0, '本月产量（吨）', sc('产量'), QTY, to=SH_PROD)
    card(ws, r + 1, 1, '本月销量（吨，产品）', sc('销量'), QTY, to=SH_SALE)
    card(ws, r + 1, 2, '成品库存（吨）', sc('成品'), QTY, big=True, to=SH_STK)
    ws.row_dimensions[r + 1].height = 26
    hr = R_MH
    header(ws, hr, [('B', '主原料'), ('C', '现在库存（吨）'), ('D', '近 30 天日均投料\n（吨/天）'), ('F', '够用几天'), ('G', '说明')],
           C_VIEW, height=32)
    put(ws, f'H{hr}', None, None, fill(C_VIEW))
    ws.merge_cells(f'G{hr}:{LAST}{hr}')
    link(ws[f'B{hr}'], SH_STK)
    R_ = RM
    for i in range(1, SHOW_RM + 1):
        rr = hr + i
        k = f'${R_["显示"]}${i + 1}'
        g = lambda f: f'INDEX({rng(R_[f], N_RM)},{k})'
        put(ws, f'B{rr}', f'=IF({k}=0,"",{g("名称")})', F_TXTB, LBL, align=AL)
        put(ws, f'C{rr}', f'=IF({k}=0,"",{g("库存")})', F_AUTOB, WHITE, QTY, AR)
        put(ws, f'D{rr}', f'=IF({k}=0,"",{g("日均")})', F_AUTOB, WHITE, QTY, AR)
        put(ws, f'F{rr}', f'=IF({k}=0,"",{g("天")})', F_TXTB, WHITE, DAYS, AR)
        for cc in 'GHIJKL':
            put(ws, f'{cc}{rr}', None, border=False)
        ws.row_dimensions[rr].height = 20
    put(ws, f'G{hr + 1}', (f'="主原料＝物料档案里类别是「原料」的，有库存或近 30 天投过料的排前面，最多 {SHOW_RM} 个。'
                           f'库存＝截止日那个月月初账面＋本月到截止日的进货、废料回收－投料、卖出＋月末盘点差。'
                           f'日均投料＝"&TEXT(P_截止-{sc("天数")}+1,"m/d")&"～"&TEXT(P_截止,"m/d")&" 共 "&{sc("天数")}&" 天的投料÷天数；'
                           f'够用几天＝库存÷日均（不足一天算 0），没投料的显示「—」。"'), F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'G{hr + 1}:{LAST}{hr + SHOW_RM}')
    ws.conditional_formatting.add(f'C{hr + 1}:C{hr + SHOW_RM}', FormulaRule(formula=[f'AND(ISNUMBER(C{hr + 1}),C{hr + 1}<-0.0005)'],
                                                                            font=F_RED))
    ws.conditional_formatting.add(f'F{hr + 1}:F{hr + SHOW_RM}', FormulaRule(formula=[f'AND(ISNUMBER(F{hr + 1}),F{hr + 1}<7)'],
                                                                            font=F_RED))
    ws.conditional_formatting.add(f'B{hr + 1}:F{hr + SHOW_RM}', FormulaRule(formula=[f'$B{hr + 1}=""'], fill=cf_bg('FFFFFFFF')))

    # ── ④ 成本利润 ──
    r = R4
    section(ws, r, 'B', LAST, '④ 成本、利润（按记账凭证，一个月一个月算；点蓝字看利润表、成本费用分析）', C_RPT)
    th = R_TH
    m, m1 = sc('m'), sc('m1')
    header(ws, th, [('B', '项目'),
                    ('C', f'=IF({sc("ok1")}=1,{m1}&" 月（上个完整月）","上个月（建账以前，没有）")'),
                    ('D', f'={m}&" 月"&IF({sc("整月")}=1,"（整月）","（截至 "&TEXT(P_截止,"m/d")&"，还没过完）")'),
                    ('F', f'="本年（"&P_建账月&"～"&{m}&" 月）"'), ('G', '口径')], C_RPT, height=32)
    put(ws, f'H{th}', None, None, fill(C_RPT))
    ws.merge_cells(f'G{th}:{LAST}{th}')
    mo = lambda k, j: f'INDEX(${MO[k]}$2:${MO[k]}$13,{j})'
    pl = lambda nm, j: f'INDEX({nm},{j})'

    def per_month(key, j):
        if key == '制造':
            return f'IF({mo("产量", j)}>0,ROUND({mo("完工", j)}/{mo("产量", j)},2),"—")'
        if key == '完全':
            return (f'IF(AND({mo("产量", j)}>0,{mo("销量", j)}>0),ROUND(ROUND({mo("完工", j)}/{mo("产量", j)},2)'
                    f'+{mo("费运", j)}/{mo("销量", j)},2),"—")')
        if key == '均价':
            return f'IF(ISNUMBER({pl("利润表_均价", j)}),{pl("利润表_均价", j)},"—")'
        if key == '毛利':
            return f'IF(ISNUMBER({pl("利润表_每吨毛利", j)}),{pl("利润表_每吨毛利", j)},"—")'
        return mo({'产量': '产量', '销量': '销量', '净利': '净利', '分后': '分后'}[key], j)

    def year(key):
        S_ = lambda k: sc('S' + k)
        if key == '制造':
            return f'IF({S_("产")}>0,ROUND({S_("完")}/{S_("产")},2),"—")'
        if key == '完全':
            return f'IF(AND({S_("产")}>0,{S_("销")}>0),ROUND(ROUND({S_("完")}/{S_("产")},2)+{S_("费")}/{S_("销")},2),"—")'
        if key == '均价':
            return f'IF({S_("销")}>0,ROUND({S_("收")}/{S_("销")},2),"—")'
        if key == '毛利':
            return f'IF({S_("销")}>0,ROUND(ROUND({S_("收")}/{S_("销")},2)-ROUND({S_("成")}/{S_("销")},2),2),"—")'
        return S_({'产量': '产', '销量': '销', '净利': '净', '分后': '分'}[key])

    lines4 = [
        ('每吨制造成本', '制造', MONEY, '＝完工入库成本（直接材料＋分到的制造费用，所有产品）÷ 产量', SH_COST),
        ('每吨完全成本', '完全', MONEY, '＝每吨制造成本 ＋（销售费用、管理费用、财务费用 ＋ 销货运费）÷ 销量', SH_COST),
        ('不含税均价', '均价', MONEY, '＝主营业务收入（不含税）÷ 销量（利润表）', SH_PL),
        ('每吨毛利', '毛利', MONEY,
         '="＝均价 － 每吨主营业务成本"&IF(P_运费进成本=1,"（含销货运费）","（销货运费在销售费用里）")&"（利润表）"', SH_PL),
        ('产量（吨）', '产量', QTY, '生产日报的成品吨（整月，所有产品）', SH_PROD),
        ('销量（吨）', '销量', QTY, '销售台账的产品吨（整月，不含卖料）', SH_SALE),
        ('利润', '净利', MONEY, '利润表的净利润（按记账凭证；截止日以后月末的折旧、摊销、利息、成本结转还没做）', SH_PL),
        ('分完投资利息后的利润', '分后', MONEY, '＝利润 － 合伙人投资款按约定算的「利息」（利润分配，不进费用）', SH_PL),
    ]
    for i, (lab, key, fmt, note, to) in enumerate(lines4):
        rr = th + 1 + i
        bold = key in ('完全', '毛利', '分后')
        fl = FILL_TOT if key == '分后' else WHITE
        link(put(ws, f'B{rr}', lab, LINKF, LBL if key != '分后' else FILL_TOT, align=AL), to)
        put(ws, f'C{rr}', f'=IF({sc("ok1")}=1,{per_month(key, m1)},"—")', F_TXTB if bold else F_AUTOB, fl, fmt, AR)
        put(ws, f'D{rr}', f'={per_month(key, m)}', F_TXTB if bold else F_AUTOB, fl, fmt, AR)
        put(ws, f'F{rr}', f'={year(key)}', F_TXTB if bold else F_AUTOB, fl, fmt, AR)
        put(ws, f'G{rr}', note, F_NOTE, align=ALW)
        for cc in 'HIJKL':
            put(ws, f'{cc}{rr}', None)
        ws.merge_cells(f'G{rr}:{LAST}{rr}')
        ws.row_dimensions[rr].height = 22
    assert len(lines4) == N4
    end4 = th + len(lines4)
    ws.conditional_formatting.add(f'C{th + 1}:F{end4}', FormulaRule(formula=[f'AND(ISNUMBER(C{th + 1}),C{th + 1}<-0.005)'], font=F_RED))
    put(ws, f'B{end4 + 1}', ('每吨数是全部产品合在一起的；要看某一个产品的单耗、收率、每吨成本拆项，去【成本费用分析】选产品。'
                            '「本年」的每吨数用本年合计重新除，不是各月相加。'), F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'B{end4 + 1}:{LAST}{end4 + 1}')
    ws.row_dimensions[end4 + 1].height = 26

    # ── ⑤ 提醒 ──
    r = R5
    assert r == end4 + 3
    section(ws, r, 'B', LAST, '⑤ 提醒（点那一行跳过去）', C_CHK)
    rem = [
        (f'=IF(N(校验_错数)>0,"✗ 数据校验有 "&N(校验_错数)&" 项要改","√ 数据校验没有要改的 ✗")'
         f'&IF(N(校验_警数)>0,"，"&N(校验_警数)&" 项要看看（⚠）","")&"  → 点这里看【数据校验】"', SH_CHK),
        ('=IF(N(校验_待补票笔数)>0,"⚠ 还有 "&N(校验_待补票笔数)&" 笔待补票（"&TEXT(N(校验_待补票金额),"#,##0.00")&" 元）：'
         '补到票后在【发票登记】登一行，垫付编号填那一笔的编号","√ 待补票的都补上了")', SH_ADV),
        ('=IF(N(校验_待定价进货笔数)>0,"⚠ 还有 "&N(校验_待定价进货笔数)&" 车进货没定价（按参考单价暂估 "'
         '&TEXT(N(校验_待定价进货金额),"#,##0.00")&" 元）：定了价回【进货物流登记】原行补单价","√ 进货都定价了")', SH_PUR),
        (f'=IF({sc("待销数")}>0,"⚠ 还有 "&{sc("待销数")}&" 车发货没定价（"&TEXT({sc("待销吨")},"#,##0.000")&" 吨，收入还没算进去）：'
         f'定了价回【销售台账】原行补单价","√ 发货都定价了")', SH_SALE),
        ('=IF(N(校验_负库存)>0,"✗ 有 "&N(校验_负库存)&" 种原料、产品月末账面数量是负数：多半漏登了进货、产量，或者投料、卖出登多了'
         '（看【进销存】）","√ 没有负库存")', SH_STK),
        ('=IF(N(校验_工资没登)>0,"⚠ 有 "&N(校验_工资没登)&" 个月发了工资却没登工资表（月底一人一行，自动计提工资）",'
         '"√ 发了工资的月份工资表都登了")', SH_WAGE),
        ('=IF(N(校验_固定支出没付条数)>0,"⚠ 本月固定支出有 "&N(校验_固定支出没付条数)&" 条过了日子还没付清（还差 "'
         '&TEXT(N(校验_固定支出没付金额),"#,##0.00")&" 元）：付了以后资金台帐「编号」填固定支出的编号","√ 本月到期的固定支出都付了")',
         SH_FIX),
        (f'=IF({sc("示例")}>0,"⚠ 还有 "&{sc("示例")}&" 行示例没删（灰底、备注写着「示例」）：正式用之前整行删掉","√ 示例行都删了")', SH_CASH),
        ('=IF(P_公司名称="","⚠ 【基础资料】还没填公司名称（首页、对账单、凭证抬头要用）","√ 公司名称："&P_公司名称)', SH_BASE),
    ]
    for i, (f, to) in enumerate(rem):
        line(ws, r + 1 + i, f, to, h=20, font=F_TXTB if i == 0 else F_TXT)
    rr0, rr1 = r + 1, r + len(rem)
    ws.conditional_formatting.add(f'B{rr0}:B{rr1}', FormulaRule(formula=[f'OR(LEFT(B{rr0},1)="✗",LEFT(B{rr0},1)="⚠")'],
                                                                font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'B{rr0}:B{rr1}', FormulaRule(formula=[f'LEFT(B{rr0},1)="√"'], font=F_OK))

    # ── ⑥ 导航 ──
    r = rr1 + 2
    section(ws, r, 'B', LAST, '⑥ 各张表（点名字跳过去；每张表右上角有「← 回首页」）', C_HOME)
    r += 1
    col_sets = [('B', 'C', 'D'), ('F', 'G', 'H'), ('J', 'K', 'L')]
    acc_grp = None
    for grp, color, sheets in GROUPS:
        if grp == '首页':
            continue
        if grp == '会计':
            acc_grp = (grp, color, sheets)
            continue
        r = _nav_block(ws, r, GRP_NAME.get(grp, grp), color, sheets, col_sets)
    if acc_grp:
        r += 1
        section(ws, r, 'B', LAST, '会计用（给会计看、打印；老板平时不用看）', acc_grp[1])
        r = _nav_block(ws, r + 1, None, acc_grp[1], acc_grp[2], col_sets)

    # ── 启用三步 ──
    r += 1
    section(ws, r, 'B', LAST, '启用三步（第一次用之前做一遍）', C_HOME)
    for i, t in enumerate(STEPS):
        rr = r + 1 + i
        put(ws, f'B{rr}', t, F_TXT, align=ALW, border=False)
        ws.merge_cells(f'B{rr}:{LAST}{rr}')
        ws.row_dimensions[rr].height = 46 if i == 1 else 32
    r = r + len(STEPS) + 2

    # ── 常见情况怎么记 ──
    section(ws, r, 'B', LAST, '常见情况怎么记', C_HOME)
    for i, (kk, v) in enumerate(HOWTO):
        rr = r + 1 + i
        put(ws, f'B{rr}', kk, F_TXTB, GREY, align=ALW)
        put(ws, f'C{rr}', v, F_TXT, align=ALW)
        for cc in 'DEFGHIJKL':
            put(ws, f'{cc}{rr}', None)
        ws.merge_cells(f'C{rr}:{LAST}{rr}')
        ws.row_dimensions[rr].height = 32 if len(v) < 90 else 44
    r = r + len(HOWTO) + 2

    # ── 注意 ──
    section(ws, r, 'B', LAST, '注意', C_HOME)
    nt = notes()
    for i, t in enumerate(nt):
        rr = r + 1 + i
        put(ws, f'B{rr}', t, F_TXT, align=ALW, border=False)
        ws.merge_cells(f'B{rr}:{LAST}{rr}')
        ws.row_dimensions[rr].height = 30
    end = r + len(nt)

    hide(ws, *[CL(c) for c in range(CI(SCL), CI(PD['库存']) + 1)])
    print_setup(ws, None, landscape=False)
    ws.print_area = f'A1:{LAST}{end}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    ws.sheet_view.showGridLines = False


def _nav_block(ws, r, name, color, sheets, col_sets):
    if name:
        put(ws, f'B{r}', name, F_SEC, fill(color), align=AL)
        ws.merge_cells(f'B{r}:{LAST}{r}')
        r += 1
    for j, sh in enumerate(sheets):
        c1, c2, c3 = col_sets[j % 3]
        rr = r + j // 3
        link(put(ws, f'{c1}{rr}', sh, LINKF, align=AL), sh)
        put(ws, f'{c2}{rr}', DESC.get(sh, ''), F_NOTE, align=ALW)
        ws.merge_cells(f'{c2}{rr}:{c3}{rr}')
        ws.row_dimensions[rr].height = 28
    return r + (len(sheets) + 2) // 3
