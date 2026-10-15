# -*- coding: utf-8 -*-
"""A075 矿业美元财务账套 —— 生成脚本入口。10/15 起出「录入本＋报表本」两本联动（用户：录入、报表分两本，录的时候不卡）。

    python3 -I build.py [输出目录] [--one 一本版.xlsx]     只生成两本（openpyxl 写盘，不压共享公式、不重算）；--one 另存一本版
    python3 -I make_all.py [输出目录] [--keep]              一条龙（见 make_all.py）

两本都从同一本「一本版」来（make_onebook，在内存里；一本版＝以前那本提速版＋这次的改动），split.py 拆：
  报表本 A075-2_报表本.xlsx：一本版原样，录入本那些表变成同名同位置的隐藏取数表（外部链接指向同一文件夹的录入本）；
  录入本 A075-1_录入本.xlsx：录入表＋基础资料等＋新首页 / 录入检查 / 使用说明，只留「轻」公式。
一本版怎么来的（详见 README.md）：
  A062 生成脚本（9/26）＋ 用户 10/05 改的结构 ＋ p1008.py（A062 10/08 修改脚本）＋ A075 的改动（common.py「A075 参数」）
  ＋ 用户的东西从原件读（ref_a075.py）＋ 提速（p_opt）＋ WPS 存盘修正（p_wps）＋ 10/15 手工登记表的调整、新报表（s_extra）。
"""
import importlib.util
import os
import sys
import time

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from openpyxl import Workbook
from openpyxl.workbook.properties import CalcProperties

from common import *
import ref_a075
import s_base, s_cash, s_mods, s_mend, s_je, s_ledger, s_reports, s_check
import p1008
import p_a075
import p_opt
import p_wps
import split


def _opt(name):
    """有这个模块就导入（别人还没交的新文件：s_extra 新报表、entry 录入本写法）。模块里自己的错照常报出来"""
    return importlib.import_module(name) if importlib.util.find_spec(name) else None


ROOT = os.path.dirname(HERE)
OUT_DIR = ROOT                                   # 两本默认放 A075 文件夹（同一个文件夹，外部链接写相对文件名）
OUT = os.path.join(ROOT, '_calc', 'A075_一本版.xlsx')     # build()（只出一本版，测试 / 对照用）的默认输出


def prepare(ref):
    d = ref_a075.read(ref)
    p = d['params']
    ctx = dict(ledger=d['ledger'], params=p, fx=d['fx'], chart=d['chart'], overrides={}, counterparties=d['counterparties'],
               projects=d['projects'], rules=d['rules'], months=d['months'], accounts=d['accounts'], rates=d['rates'],
               mat_cats=d['fx'].MAT_CATS, home_year=None, home_month=None, n_fee=0, ref=ref,
               gl_default=ref[SH_GL]['B3'].value)
    for k, v in d.items():                      # 别人在 ref_a075.read 里新加的东西（手工登记表、u4 流水…）也带进 ctx
        ctx.setdefault(k, v)
    return ctx


def _order(wb):
    """表顺序：SHEET_ORDER 里的按它排；不在里面的（新表）先接在后面——别的模块新建的表不会因为不在 SHEET_ORDER 里被丢掉"""
    known = [wb[n] for n in SHEET_ORDER if n in wb.sheetnames]
    wb._sheets = known + [ws for ws in wb._sheets if ws.title not in SHEET_ORDER]


def _final_order(wb):
    """最后排一次：SHEET_ORDER 里的按它排（后建的表写进 SHEET_ORDER 也照排）；不在里面的：录入本那边的新表（货柜登记、录入检查，
    SPLIT 里有）接在手工凭证后面，别的看得见的新表（s_extra 的报表）放【数据校验】前面，隐藏的放最后"""
    names = [ws.title for ws in wb._sheets]
    rest = [n for n in SHEET_ORDER if n in names]
    new = [n for n in names if n not in SHEET_ORDER]
    ent = [n for n in new if n in SPLIT]
    vis = [n for n in new if n not in SPLIT and wb[n].sheet_state == 'visible']
    hid = [n for n in new if n not in SPLIT and n not in vis]
    at = rest.index(SH_MAN) + 1 if SH_MAN in rest else len(rest)
    rest[at:at] = ent
    at = rest.index(SH_CHK) if SH_CHK in rest else len(rest)
    rest[at:at] = vis
    wb._sheets = [wb[n] for n in rest + hid]


def make_onebook(ref_path=ref_a075.REF, log=print, tweak=None):
    """一本版（内存里）→ (wb, ctx)。tweak(wb, ctx)：最后再改一下（场景测试填数用）"""
    t0 = time.time()
    SEG_NAMES.clear()
    ref = ref_a075.load(ref_path)
    log(f'  读原件 {time.time() - t0:.0f} 秒')
    ctx = prepare(ref)
    wb = Workbook()
    wb.remove(wb.active)
    s_base.build_base(wb, ctx)
    s_base.build_rules(wb, ctx)
    s_base.build_rates(wb, ctx)
    s_base.build_accounts(wb, ctx)
    s_cash.build_cash(wb, ctx)
    s_mods.build_buy(wb, ctx)
    s_mods.build_issue(wb, ctx)
    s_mods.build_prod(wb, ctx)
    s_mods.build_sale(wb, ctx)
    s_mods.build_pay(wb, ctx)
    s_mods.build_fa(wb, ctx)
    s_mods.build_man(wb, ctx)
    wb.create_sheet(SH_AUX)
    s_mend.build_inv_grid(wb, ctx)
    s_mend.build_ore(wb, ctx)
    s_mend.build_mend(wb, ctx)
    s_je.build_je(wb, ctx)
    s_je.build_vindex(wb, ctx)
    s_je.build_vlist(wb, ctx)
    s_je.build_vprint(wb, ctx)
    s_ledger.build_tb(wb, ctx)
    s_ledger.build_gl(wb, ctx)
    s_ledger.build_acct(wb, ctx)
    s_ledger.build_arap(wb, ctx)
    s_ledger.build_inv(wb, ctx)
    s_reports.build_bs(wb, ctx)
    s_reports.build_pl(wb, ctx)
    s_reports.build_cf(wb, ctx)
    s_reports.build_ops(wb, ctx)
    s_check.build_check(wb, ctx)
    wb.create_sheet(SH_HOME)            # 首页、使用说明：p_a075 从原件整张照搬
    wb.create_sheet(SH_HELP)
    s_base.define_names(wb)
    _order(wb)
    log(f'  A062 生成脚本部分 {time.time() - t0:.0f} 秒')
    for line in p1008.apply(wb, ctx):
        log('   · 10/08 ' + line)
    for line in p_a075.apply(wb, ctx):
        log('   · A075 ' + line)
    for line in p_opt.apply(wb, ctx):
        log('   · 提速 ' + line)
    s_extra = _opt('s_extra')
    if s_extra is not None:             # 10/15 新报表（发生情况报表、记账底稿、物资出入库台账、黄金总记录、收支分类报表）
        for line in (s_extra.build(wb, ctx) or []):
            log('   · 新报表 ' + line)
    for line in p_wps.apply(wb):
        log('   · WPS ' + line)
    entry = _opt('entry')
    if entry is not None:               # 【录入检查】：一本版里也建（报表本【数据校验】第 41 项取它；报表本里它是取数表，取录入本算的结果）
        entry.build_entry_check(wb, ctx)
        if hasattr(s_check, 'report_check41') and not any('录入本：表头没挪' in str(c.value or '') for c in wb[SH_CHK]['B']):
            for line in s_check.report_check41(wb):
                log('   · 拆本 ' + line)
        if hasattr(s_check, 'home_link_line') and not any(
                isinstance(c.value, str) and '数据来自录入本' in c.value for row in wb[SH_HOME].iter_rows() for c in row):
            log('   · 拆本 报表本首页加「数据来自录入本：流水 N 笔，最后一笔…」：' + s_check.home_link_line(wb[SH_HOME]))
    wb[SH_AUX].sheet_state = 'hidden'
    wb[SH_VIDX].sheet_state = 'hidden'
    wb[SH_CLOSE].sheet_state = 'hidden'
    if entry is not None and entry.ECHK in wb.sheetnames:
        wb[entry.ECHK].sheet_state = 'hidden'
    _final_order(wb)
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SH_HOME
    if tweak:
        tweak(wb, ctx)
    log(f'  一本版（内存）{time.time() - t0:.0f} 秒')
    return wb, ctx


def _save(wb, path, log=print):
    t = time.time()
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=False)     # 不写 fullCalcOnLoad（打开整本重算太慢；缓存值由 make_all 写回）
    wb.save(path)
    log(f'✓ 存 {os.path.basename(path)}（{time.time() - t:.0f} 秒）')


def build(out=OUT, ref_path=ref_a075.REF, log=print, tweak=None):
    """只出一本版（以前的成品形态；对照 / 场景测试用）"""
    wb, ctx = make_onebook(ref_path, log, tweak)
    _save(wb, out, log)
    return ctx


def build_split(out_dir=OUT_DIR, ref_path=ref_a075.REF, log=print, tweak=None, onebook=None):
    """出两本（＋可选一本版）→ (录入本路径, 报表本路径, ctx)。都在同一个内存 wb 上：一本版 → 存一本版 → 改成报表本、存、退回 → 拆成录入本、存"""
    t0 = time.time()
    out1, out2 = os.path.join(out_dir, WB1_FILE), os.path.join(out_dir, WB2_FILE)
    wb, ctx = make_onebook(ref_path, log, tweak)
    if onebook:
        _save(wb, onebook, log)
    undo = split.make_report(wb, ctx, log)
    _save(wb, out2, log)
    undo()
    split.make_entry(wb, ctx, log)
    _save(wb, out1, log)
    log(f'✓ 两本生成完（{time.time() - t0:.0f} 秒）')
    return out1, out2, ctx


if __name__ == '__main__':
    args = [a for a in sys.argv[1:] if not a.startswith('--')]
    one = None
    if '--one' in sys.argv:
        one = sys.argv[sys.argv.index('--one') + 1]
        args = [a for a in args if a != one]
    if args and args[0].lower().endswith('.xlsx'):
        sys.exit('参数是输出目录（两本都放进去），不是 .xlsx 文件名；只要一本版用 --one 一本版.xlsx')
    build_split(args[0] if args else OUT_DIR, onebook=one)
