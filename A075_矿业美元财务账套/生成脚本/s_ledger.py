# -*- coding: utf-8 -*-
"""账簿：【科目余额表】【明细账】【资金账户余额表】【往来余额表】【存货收发存】。
全部从【记账分录】取数（期间跟【首页】的报表年度/月份走），科目编码是文本，上级科目用「编码*」通配汇总下级。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from openpyxl.workbook.defined_name import DefinedName

from common import *
from s_base import code_of
from s_mend import inv_field, INV_R0, INV_R1

TB_R1 = TB_R0 + (AC_R1 - AC_R0)


def _dn(wb, name, ref):
    wb.defined_names[name] = DefinedName(name, attr_text=ref)


def tb_agg(col, r):
    """科目余额表第 r 行：末级＝本行自己（col 列）；上级＝本表里编码以它开头的各行自己的数之和（Y＝1：重复编码只算第一行）"""
    return (f'IF(${T_LEAF}{r}="是",${col}{r},SUMIFS(${col}${TB_R0}:${col}${TB_R1},${T_CODE}${TB_R0}:${T_CODE}${TB_R1},'
            f'${T_CODE}{r}&"*",$Y${TB_R0}:$Y${TB_R1},1))')


def je_sum(side, code_crit, *date_crit):
    """SUMIFS(分录借方/贷方, 分录科目, 条件, 分录日期, 条件…)"""
    col = '分录借方' if side == 'D' else '分录贷方'
    extra = ''.join(f',分录日期,{d}' for d in date_crit)
    return f'SUMIFS({col},分录科目,{code_crit}{extra})'


def je_net(code_crit, *date_crit):
    return f'({je_sum("D", code_crit, *date_crit)}-{je_sum("C", code_crit, *date_crit)})'


def _kpi_band(ws, last_col, row=3):
    ws.row_dimensions[row].height = 26
    for i in range(1, CI(last_col) + 1):
        put(ws, f'{CL(i)}{row}', None, fill_=FILL_KPI, border=False)


def _kpi(ws, lab_cell, lab, val_cell, f, fmt=None, merge_to=None):
    put(ws, lab_cell, lab, F_KPI_L, FILL_KPI, align=AR_, border=False)
    put(ws, val_cell, f, Font(name=YH, sz=12, bold=True, color='FFC00000'), FILL_KPI, fmt, AL, border=False)
    if merge_to:
        ws.merge_cells(f'{val_cell}:{merge_to}')


# ─────────────────────────── 科目余额表 ───────────────────────────
def build_tb(wb, ctx):
    ws = wb.create_sheet(SH_TB)
    widths(ws, {T_CODE: 11, T_NAME: 34, T_DIR: 5, T_LVL: 5, T_LEAF: 5, T_OB_D: 17, T_OB_C: 17, T_M_D: 16, T_M_C: 16,
                T_Y_D: 17, T_Y_C: 17, T_CB_D: 17, T_CB_C: 17, T_YNET: 16, T_ONET: 16, T_CNET: 16, T_POS: 6})
    title(ws, '科目余额表（按【首页】报表年度 / 月份）', T_CB_C, C_LED,
          '💡 全部从【记账分录】汇总：月初余额、本月发生、本年累计、月末余额；上级科目＝下级合计。余额按借贷方向拆开列（借方余额列借方，贷方余额列贷方）。'
          '损益类科目不结转（表结法），本年累计就是利润表的数；未分配利润在资产负债表里自动算。最下面一行是一级科目合计，借贷应该相等。')
    _kpi_band(ws, T_CB_C)
    _kpi(ws, 'A3', '期间：', 'B3', f'={REP_M0}', MONTH)
    lvl1 = lambda col: f'SUMIF(${T_LVL}${TB_R0}:${T_LVL}${TB_R1},1,{col}{TB_R0}:{col}{TB_R1})'
    chk = (f'=IF(AND(ROUND({lvl1(T_OB_D)}-{lvl1(T_OB_C)},2)=0,ROUND({lvl1(T_M_D)}-{lvl1(T_M_C)},2)=0,'
           f'ROUND({lvl1(T_CB_D)}-{lvl1(T_CB_C)},2)=0,COUNTIFS(分录有额,1,分录全称,"（科目表里没有）")=0),"√ 借贷平衡",'
           f'IF(COUNTIFS(分录有额,1,分录全称,"（科目表里没有）")>0,"✗ 有分录的科目不在科目表里（去【记账分录】筛「科目表里没有」）",'
           f'"✗ 借贷不平，差 "&FIXED({lvl1(T_CB_D)}-{lvl1(T_CB_C)},2)))')
    _kpi(ws, 'D3', '平衡检查：', 'F3', chk, merge_to='J3')
    ws.merge_cells('D3:E3')
    put(ws, 'A4', '月初＝本月第一天之前；本年累计＝1 月 1 日到月末。', F_NOTE, border=False)
    # 两层表头
    for c0, c1, t in ((T_OB_D, T_OB_C, '月初余额'), (T_M_D, T_M_C, '本月发生额'), (T_Y_D, T_Y_C, '本年累计发生额'),
                      (T_CB_D, T_CB_C, '月末余额')):
        ws.merge_cells(f'{c0}{TB_HDR - 1}:{c1}{TB_HDR - 1}')
        put(ws, f'{c0}{TB_HDR - 1}', t, F_HDR, fill(C_LED), align=AC)
        put(ws, f'{c1}{TB_HDR - 1}', None, fill_=fill(C_LED))
    for c in (T_CODE, T_NAME, T_DIR, T_LVL, T_LEAF):
        put(ws, f'{c}{TB_HDR - 1}', None, fill_=fill(C_LED))
    header(ws, TB_HDR, [(T_CODE, '科目编码'), (T_NAME, '科目名称'), (T_DIR, '方向'), (T_LVL, '级次'), (T_LEAF, '末级'),
                        (T_OB_D, '借方'), (T_OB_C, '贷方'), (T_M_D, '借方'), (T_M_C, '贷方'), (T_Y_D, '借方'),
                        (T_Y_C, '贷方'), (T_CB_D, '借方'), (T_CB_C, '贷方'), (T_YNET, '年初净额'), (T_ONET, '月初净额'),
                        (T_CNET, '月末净额'), (T_POS, '科目行'), ('R', '期末直接'), ('S', '年初直接')], C_LED, height=24)
    for i in range(TB_R1 - TB_R0 + 1):
        r = TB_R0 + i
        k = i + 1
        P = f'${T_POS}{r}'
        A = f'${T_CODE}{r}'
        ws[f'{T_POS}{r}'] = f'=IFERROR(MATCH(SMALL({acc(AC_KEY)},{k}),{acc(AC_KEY)},0),"")'
        ws[f'{T_CODE}{r}'] = f'=IF({P}="","",INDEX({acc(AC_CODE)},{P})&"")'
        ws[f'{T_NAME}{r}'] = f'=IF({P}="","",REPT("　",N(INDEX({acc(AC_LVL)},{P}))-1)&INDEX({acc(AC_NAME)},{P}))'
        ws[f'{T_DIR}{r}'] = f'=IF({P}="","",INDEX({acc(AC_DIR)},{P})&"")'
        ws[f'{T_LVL}{r}'] = f'=IF({P}="","",INDEX({acc(AC_LVL)},{P}))'
        ws[f'{T_LEAF}{r}'] = f'=IF({P}="","",INDEX({acc(AC_LEAF)},{P})&"")'
        crit = f'{A}&"*"'
        opn = f'SUMIF({acc(AC_CODE)},{crit},{acc(AC_OPEN)})'
        # opt：原来每格 SUMIFS(分录借方,分录科目,编码&"*",分录日期,…) 在 1.8 万行上做通配符匹配（342 个科目 × 8 列）。
        # 现在：记账分录 U 列「余额键」＝科目行×10000＋月序；这里先算「本科目自己」的几个数（隐藏 Z～AD，只有 X＝有分录的才算，
        # 精确匹配数字键），末级科目直接用，上级科目＝本表里「编码*」各行自己的数加起来（只扫本表 450 行，重复编码只算一次），
        # 科目表里没有 / 进不了本表的编码（余额键是 "#编码"）照原来按编码前缀＋日期另外加（平时一笔都没有，整列不算）。
        Q = f'${T_POS}{r}'
        base = f'{Q}*{TB_KEYM}'
        ws[f'X{r}'] = f'=IF({Q}="",0,IF(ISNUMBER(MATCH({Q},分录科目行,0)),1,0))'
        ws[f'Y{r}'] = f'=IF({Q}="",0,1)' if r == TB_R0 else f'=IF({Q}="",0,IF({Q}={T_POS}{r - 1},0,1))'
        ws[f'Z{r}'] = f'=IF($X{r}=0,0,SUMIF(分录余额键,{base}+$X$3,分录借方))'
        ws[f'AA{r}'] = f'=IF($X{r}=0,0,SUMIF(分录余额键,{base}+$X$3,分录贷方))'
        yr = f'分录余额键,">"&({base}+$X$4+0.5),分录余额键,"<="&({base}+$X$3)'
        ws[f'AB{r}'] = f'=IF($X{r}=0,0,SUMIFS(分录借方,{yr}))'
        ws[f'AC{r}'] = f'=IF($X{r}=0,0,SUMIFS(分录贷方,{yr}))'
        y0 = f'分录余额键,">"&{base},分录余额键,"<"&({base}+$X$4+1)'
        ws[f'AD{r}'] = f'=IF($X{r}=0,0,SUMIFS(分录借方,{y0})-SUMIFS(分录贷方,{y0}))'
        m = (f'">="&{REP_M0}', f'"<="&{REP_M1}')
        y = (f'">="&{REP_Y0}', f'"<="&{REP_M1}')
        osum = lambda side, *dc: (f'SUMIFS({"分录借方" if side == "D" else "分录贷方"},分录余额键,"#"&{A}&"*"'
                                  + ''.join(f',分录日期,{d}' for d in dc) + ')')
        orph = lambda side, *dc: f'IF($X$5>0,{osum(side, *dc)},0)'
        ws[f'{T_M_D}{r}'] = f'=IF({P}="","",{tb_agg("Z", r)}+{orph("D", *m)})'
        ws[f'{T_M_C}{r}'] = f'=IF({P}="","",{tb_agg("AA", r)}+{orph("C", *m)})'
        ws[f'{T_Y_D}{r}'] = f'=IF({P}="","",{tb_agg("AB", r)}+{orph("D", *y)})'
        ws[f'{T_Y_C}{r}'] = f'=IF({P}="","",{tb_agg("AC", r)}+{orph("C", *y)})'
        lt = chr(34) + "<" + chr(34) + "&" + REP_Y0
        ws[f'AE{r}'] = f'=IF({P}="",0,{tb_agg("AD", r)}+IF($X$5>0,{osum("D", lt)}-{osum("C", lt)},0))'
        ws[f'{T_YNET}{r}'] = f'=IF({P}="","",ROUND({opn}+$AE{r},2))'
        # 月初＝年初＋（本年－本月）：原来另做两次「日期＜月初」的 SUMIFS
        ws[f'{T_ONET}{r}'] = (f'=IF({P}="","",ROUND({opn}+$AE{r}+({T_Y_D}{r}-{T_M_D}{r})-({T_Y_C}{r}-{T_M_C}{r}),2))')
        ws[f'{T_CNET}{r}'] = f'=IF({P}="","",ROUND({T_ONET}{r}+{T_M_D}{r}-{T_M_C}{r},2))'
        ws[f'{T_OB_D}{r}'] = f'=IF({P}="","",MAX({T_ONET}{r},0))'
        ws[f'{T_OB_C}{r}'] = f'=IF({P}="","",MAX(-{T_ONET}{r},0))'
        ws[f'{T_CB_D}{r}'] = f'=IF({P}="","",MAX({T_CNET}{r},0))'
        ws[f'{T_CB_C}{r}'] = f'=IF({P}="","",MAX(-{T_CNET}{r},0))'
        # 直接记在本科目上的（末级＝全部；上级＝减掉下一级明细）：报表按余额方向拆应收应付时用，上级科目直接记的也拆得进去
        kids = lambda col: f'SUMIFS(${col}${TB_R0}:${col}${TB_R1},${T_CODE}${TB_R0}:${T_CODE}${TB_R1},{A}&"??")'
        ws[f'R{r}'] = f'=IF({P}="","",IF({T_LEAF}{r}="是",{T_CNET}{r},ROUND({T_CNET}{r}-{kids(T_CNET)},2)))'
        ws[f'S{r}'] = f'=IF({P}="","",IF({T_LEAF}{r}="是",{T_YNET}{r},ROUND({T_YNET}{r}-{kids(T_YNET)},2)))'
    cols = [T_CODE, T_NAME, T_DIR, T_LVL, T_LEAF, T_OB_D, T_OB_C, T_M_D, T_M_C, T_Y_D, T_Y_C, T_CB_D, T_CB_C]
    money = [T_OB_D, T_OB_C, T_M_D, T_M_C, T_Y_D, T_Y_C, T_CB_D, T_CB_C]
    style_rows(ws, TB_R0, TB_R1, cols, auto=tuple(cols), fmts={**{c: MONEY for c in money}, T_CODE: '@'},
               aligns={T_CODE: AL, T_NAME: AL, **{c: AR_ for c in money}})
    for r in range(TB_R0, TB_R1 + 1):
        for c in (T_YNET, T_ONET, T_CNET, T_POS, 'R', 'S'):
            ws[f'{c}{r}'].font = F_NOTE
            ws[f'{c}{r}'].number_format = MONEY
    hide(ws, [T_YNET, T_ONET, T_CNET, T_POS, 'R', 'S'])
    # opt 隐藏辅助列（T～W 是 p1008 加的「不含结转损益」，AF～AI 也在 p1008 里写）
    for c, t in (('X', '有分录'), ('Y', '首行'), ('Z', '本月借\n(本科目)'), ('AA', '本月贷\n(本科目)'), ('AB', '本年借\n(本科目)'),
                 ('AC', '本年贷\n(本科目)'), ('AD', '年初净额\n(本科目)'), ('AE', '年初净额\n(含下级)')):
        put(ws, f'{c}{TB_HDR}', t, F_NOTE, FILL_AUTO, align=ACW)
        ws.column_dimensions[c].width = 12
        for r in range(TB_R0, TB_R1 + 1):
            ws[f'{c}{r}'].font = F_NOTE
            ws[f'{c}{r}'].number_format = MONEY if c not in ('X', 'Y') else '0'
    put(ws, 'Y3', '报表月序', F_NOTE, border=False)
    put(ws, 'X3', f'={REP_N}', F_NOTE, border=False)
    put(ws, 'Y4', '上年末月序', F_NOTE, border=False)
    put(ws, 'X4', f'=$X$3-报表月份', F_NOTE, border=False)
    put(ws, 'Y5', '←前缀另算笔数', F_NOTE, border=False)
    put(ws, 'X5', '=COUNTIF(分录余额键,"#*")', F_NOTE, border=False)
    hide(ws, ['X', 'Y', 'Z', 'AA', 'AB', 'AC', 'AD', 'AE'])
    # 一级科目加粗
    ws.conditional_formatting.add(f'A{TB_R0}:{T_CB_C}{TB_R1}',
                                  FormulaRule(formula=[f'${T_LVL}{TB_R0}=1'], font=Font(name=YH, sz=10, bold=True, color='FF000000'),
                                              fill=fill('FFEEF3FA')))
    r = TB_R1 + 1
    put(ws, f'{T_CODE}{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'{T_NAME}{r}', '（一级科目合计，借＝贷）', F_TXTB, FILL_TOT, align=AL)
    for c in (T_DIR, T_LVL, T_LEAF):
        put(ws, f'{c}{r}', None, F_TXTB, FILL_TOT)
    for c in money:
        put(ws, f'{c}{r}', f'={lvl1(c)}', F_TXTB, FILL_TOT, MONEY, AR_)
    ws.auto_filter.ref = f'A{TB_HDR}:{T_CB_C}{TB_R1}'
    ws.freeze_panes = f'C{TB_R0}'
    lock_formula_sheet(ws)
    _dn(wb, '报表月序', f"{q(SH_TB)}!$X$3")          # opt：报表月是起始月份起第几个月（各表按「月序」精确汇总用）
    for nm, col in (('余额编码', T_CODE), ('余额级次', T_LVL), ('余额末级', T_LEAF), ('余额年初', T_YNET), ('余额月初', T_ONET),
                    ('余额期末', T_CNET), ('余额本月借', T_M_D), ('余额本月贷', T_M_C), ('余额本年借', T_Y_D),
                    ('余额本年贷', T_Y_C), ('余额期末直接', 'R'), ('余额年初直接', 'S')):
        _dn(wb, nm, rng(SH_TB, col, TB_R0, TB_R1))
    return ws


# ─────────────────────────── 明细账 ───────────────────────────
GL = dict(date='A', no='B', memo='C', acc='D', dr='E', cr='F', dir='G', bal='H', cp='I', src='J', jrow='K', net='L')
GL_R1 = GL_R0 + GL_N - 1


def build_gl(wb, ctx):
    ws = wb.create_sheet(SH_GL)
    g = GL
    widths(ws, {g['date']: 11, g['no']: 8, g['memo']: 42, g['acc']: 34, g['dr']: 17, g['cr']: 17, g['dir']: 5,
                g['bal']: 18, g['cp']: 12, g['src']: 8, g['jrow']: 7, g['net']: 16})
    title(ws, '明细账（选一个科目，看它每一笔）', g['src'], C_LED,
          '💡 在 B3 选科目（下拉，或者直接填编码；选上级科目会把下级全部列出来），起止日期默认是【首页】报表年度 1 月 1 日到报表月末，'
          '可以直接改。按日期、凭证号排好，最右边是逐笔余额。')
    ws.row_dimensions[3].height = 28
    for i in range(1, CI(g['src']) + 1):
        put(ws, f'{CL(i)}3', None, fill_=FILL_KPI, border=False)
    put(ws, 'A3', '科目：', F_KPI_L, FILL_KPI, align=AR_, border=False)
    ws.merge_cells('B3:C3')
    put(ws, 'B3', ctx.get('gl_default', '100207 NBU银行苏姆结算户-6688'), F_SEL, FILL_SEL, '@', AL)
    put(ws, 'D3', '起始日期：', F_KPI_L, FILL_KPI, align=AR_, border=False)
    put(ws, 'E3', f'={REP_Y0}', F_SEL, FILL_SEL, DATE, AC)
    put(ws, 'F3', '截止日期：', F_KPI_L, FILL_KPI, align=AR_, border=False)
    put(ws, 'G3', f'={REP_M1}', F_SEL, FILL_SEL, DATE, AC)
    ws.merge_cells('G3:H3')
    put(ws, 'L3', f'={code_of("B3")}', F_NOTE, border=False)
    add_list_dv(ws, 'B3', '=科目选择', '从下拉选，或直接填编码', stop=False)
    _dn(wb, '明细账科目', f"{q(SH_GL)}!$L$3")
    _dn(wb, '明细账起', f"{q(SH_GL)}!$E$3")
    _dn(wb, '明细账止', f"{q(SH_GL)}!$G$3")
    sel = '明细账科目'
    dirx = f'IFERROR(INDEX({acc(AC_DIR)},MATCH({sel},{acc(AC_CODE)},0)),"借")'
    put(ws, 'A4', '科目全称：', F_TXTB, align=AR_, border=False)
    put(ws, 'B4', f'=IFERROR(INDEX({acc(AC_FULL)},MATCH({sel},{acc(AC_CODE)},0)),"✗ 科目表里没有这个编码")', F_TXTB, align=AL, border=False)
    ws.merge_cells('B4:C4')
    put(ws, 'D4', f'="余额方向："&{dirx}&"　　本期借方："&FIXED(SUM({g["dr"]}{GL_R0}:{g["dr"]}{GL_R1}),2)'
                  f'&"　　本期贷方："&FIXED(SUM({g["cr"]}{GL_R0}:{g["cr"]}{GL_R1}),2)', F_TXT, align=AL, border=False)
    ws.merge_cells('D4:H4')
    put(ws, 'I4', f'=IF($N$5>{GL_N},"✗ 超过 {GL_N} 笔，只列了前 {GL_N} 笔，把日期范围缩小","")', F_RED, align=AL, border=False)
    # opt：N5＝要列的笔数（记账分录 V 列「明细累计」的最后一个数＝COUNT(分录明细键)）
    put(ws, 'N5', '=INDEX(分录明细累计,ROWS(分录明细累计))', F_NOTE, border=False)
    put(ws, 'M5', '笔数→', F_NOTE, border=False)
    header(ws, GL_R0 - 2, [(g['date'], '日期'), (g['no'], '凭证号'), (g['memo'], '摘要'), (g['acc'], '科目'),
                           (g['dr'], '借方'), (g['cr'], '贷方'), (g['dir'], '借/贷'), (g['bal'], '余额'), (g['cp'], '往来单位'),
                           (g['src'], '来源'), (g['jrow'], '分录行'), (g['net'], '净额')], C_LED, height=24)
    r0 = GL_R0 - 1
    opn = (f'SUMIF({acc(AC_CODE)},{sel}&"*",{acc(AC_OPEN)})+{je_net(sel + chr(38) + chr(34) + "*" + chr(34), chr(34) + "<" + chr(34) + "&明细账起")}')
    put(ws, f'{g["memo"]}{r0}', '期初余额', F_TXTB, FILL_SUBH, align=AC)
    put(ws, f'{g["net"]}{r0}', f'=IF({sel}="",0,ROUND({opn},2))', F_NOTE, fmt=MONEY)
    put(ws, f'{g["dir"]}{r0}', f'=IF(ROUND({g["net"]}{r0},2)=0,"平",IF({g["net"]}{r0}>0,"借","贷"))', F_TXTB, FILL_SUBH, align=AC)
    put(ws, f'{g["bal"]}{r0}', f'=ABS({g["net"]}{r0})', F_TXTB, FILL_SUBH, MONEY, AR_)
    for c in (g['date'], g['no'], g['acc'], g['dr'], g['cr'], g['cp'], g['src']):
        put(ws, f'{c}{r0}', None, F_TXTB, FILL_SUBH)
    for i in range(1, GL_N + 1):
        r = GL_R0 + i - 1
        K = f'${g["jrow"]}{r}'
        # opt：原来每行 MATCH(SMALL(分录明细键,i),分录明细键,0)——1,500 行 × 1.8 万行，选什么科目都照算。
        # 现在：N 列＝第 i 笔在记账分录第几行（V 列累计第一次到 i 的那一行，二分查找）、O 列＝它的明细账键、P 列＝它在这些键里排第几
        # （只在本表 1,500 格里比），K 列＝排第 i 的那一笔；用不到的行（i＞笔数）一律空着不算。
        # 超过 {GL_N} 笔时（I4 会报 ✗）照原来的写法从全部分录里挑最小的 {GL_N} 笔，结果跟原来一样。
        ws[f'N{r}'] = f'=IF({i}>$N$5,"",IFERROR(MATCH({i - 1},分录明细累计,1),0)+1)'
        ws[f'O{r}'] = f'=IF($N{r}="","",INDEX(分录明细键,$N{r}))'
        ws[f'P{r}'] = f'=IF($O{r}="","",COUNTIF($O${GL_R0}:$O${GL_R1},"<"&$O{r})+1)'
        ws[f'{g["jrow"]}{r}'] = (f'=IF($N$5>{GL_N},IFERROR(MATCH(SMALL(分录明细键,{i}),分录明细键,0),""),'
                                 f'IF({i}>$N$5,"",IFERROR(INDEX($N${GL_R0}:$N${GL_R1},MATCH({i},$P${GL_R0}:$P${GL_R1},0)),"")))')
        ws[f'{g["date"]}{r}'] = f'=IF({K}="","",INDEX(分录日期,{K}))'
        ws[f'{g["no"]}{r}'] = f'=IF({K}="","",INDEX(分录凭证号,{K}))'
        ws[f'{g["memo"]}{r}'] = f'=IF({K}="","",INDEX(分录摘要,{K}))'
        ws[f'{g["acc"]}{r}'] = f'=IF({K}="","",INDEX(分录科目,{K})&" "&INDEX(分录全称,{K}))'
        ws[f'{g["dr"]}{r}'] = f'=IF({K}="","",IF(INDEX(分录借方,{K})=0,"",INDEX(分录借方,{K})))'
        ws[f'{g["cr"]}{r}'] = f'=IF({K}="","",IF(INDEX(分录贷方,{K})=0,"",INDEX(分录贷方,{K})))'
        ws[f'{g["net"]}{r}'] = f'=IF({K}="","",{g["net"]}{r - 1}+N({g["dr"]}{r})-N({g["cr"]}{r}))' if i > 1 else \
            f'=IF({K}="","",{g["net"]}{r0}+N({g["dr"]}{r})-N({g["cr"]}{r}))'
        ws[f'{g["dir"]}{r}'] = f'=IF({K}="","",IF(ROUND({g["net"]}{r},2)=0,"平",IF({g["net"]}{r}>0,"借","贷")))'
        ws[f'{g["bal"]}{r}'] = f'=IF({K}="","",ABS({g["net"]}{r}))'
        ws[f'{g["cp"]}{r}'] = f'=IF({K}="","",INDEX(分录往来,{K}))'
        ws[f'{g["src"]}{r}'] = f'=IF({K}="","",INDEX(分录来源,{K}))'
    cols = [g[x] for x in ('date', 'no', 'memo', 'acc', 'dr', 'cr', 'dir', 'bal', 'cp', 'src')]
    style_rows(ws, GL_R0, GL_R1, cols, auto=tuple(cols),
               fmts={g['date']: DATE, g['no']: '"记-"000', g['dr']: MONEY, g['cr']: MONEY, g['bal']: MONEY},
               aligns={g['memo']: AL, g['acc']: AL, g['dr']: AR_, g['cr']: AR_, g['bal']: AR_})
    for r in range(GL_R0, GL_R1 + 1):
        ws[f'{g["jrow"]}{r}'].font = F_NOTE
        ws[f'{g["net"]}{r}'].font = F_NOTE
        ws[f'{g["net"]}{r}'].number_format = MONEY
    for c, t in (('N', '分录行'), ('O', '明细账键'), ('P', '名次')):
        put(ws, f'{c}{GL_R0 - 2}', t, F_NOTE, FILL_AUTO, align=AC)
        ws.column_dimensions[c].width = 9 if c != 'O' else 14
        for r in range(GL_R0, GL_R1 + 1):
            ws[f'{c}{r}'].font = F_NOTE
    hide(ws, [g['jrow'], g['net'], 'L', 'M', 'N', 'O', 'P'])
    ws.freeze_panes = f'C{GL_R0}'
    return ws


# ─────────────────────────── 资金账户余额表 ───────────────────────────
def build_acct(wb, ctx):
    ws = wb.create_sheet(SH_ACCT)
    cols = 'A B C D E F G H I J K L M'.split()
    widths(ws, dict(zip(cols, [24, 6, 10, 17, 17, 17, 17, 10, 18, 18, 14, 17, 17])))
    title(ws, '资金账户余额表（原币 · 折苏姆 · 跟账上核对）', 'M', C_LED,
          '💡 按【首页】报表月：每个账户月初、本月收支、月末原币余额（取【现金流水总表】），按当月记账汇率折苏姆，'
          '跟科目余额表里这个账户科目的余额核对（外币账户月末调汇以后应该正好相等；有差额多半是流水里有 ✗ 的行没记账）。')
    header(ws, 4, [('A', '账户'), ('B', '币种'), ('C', '科目'), ('D', '月初余额\n(原币)'), ('E', '本月收入\n(原币)'),
                   ('F', '本月支出\n(原币)'), ('G', '月末余额\n(原币)'), ('H', '月末\n记账汇率'), ('I', '月末折苏姆\n(按汇率)'),
                   ('J', '账上苏姆\n(科目余额表)'), ('K', '差额'), ('L', '本年收入\n(原币)'), ('M', '本年支出\n(原币)')],
           C_LED, height=40)
    ab = lambda col: rng(SH_BASE, col, ACC_R0, ACC_R1)
    R0, R1 = 5, 5 + ACC_R1 - ACC_R0
    for i in range(ACC_R1 - ACC_R0 + 1):
        r = R0 + i
        k = i + 1
        A = f'$A{r}'
        ws[f'A{r}'] = f'=IF({at(ab(A_NAME), k)}="","",{at(ab(A_NAME), k)}&"")'
        ws[f'B{r}'] = f'=IF({A}="","",{at(ab(A_CUR), k)}&"")'
        ws[f'C{r}'] = f'=IF({A}="","",{at(ab(A_CODE), k)}&"")'
        net = lambda *crit: f'SUMIFS({cash_u(K_NET)},{cash_u(K_ACC)},{A}' + ''.join(f',{cash_u(K_DATE)},{c}' for c in crit) + ')'
        side = lambda col, *crit: f'SUMIFS({cash_u(col)},{cash_u(K_ACC)},{A}' + ''.join(f',{cash_u(K_DATE)},{c}' for c in crit) + ')'
        m = (f'">="&{REP_M0}', f'"<="&{REP_M1}')
        y = (f'">="&{REP_Y0}', f'"<="&{REP_M1}')
        ws[f'D{r}'] = f'=IF({A}="","",N({at(ab(A_OPEN), k)})+{net(chr(34) + "<" + chr(34) + "&" + REP_M0)})'
        ws[f'E{r}'] = f'=IF({A}="","",{side(K_IN, *m)})'
        ws[f'F{r}'] = f'=IF({A}="","",{side(K_OUT, *m)})'
        ws[f'G{r}'] = f'=IF({A}="","",ROUND(D{r}+E{r}-F{r},2))'
        if RATE_PER_BASE:
            # A075：H 显示「1 美元 = X 原币」（直接取【月度汇率】B:I），I＝原币 ÷ 汇率
            rt = f'INDEX({q(SH_RATE)}!$B${RATE_R0}:${CL(1 + len(RATE_CURS))}${RATE_R1},{REP_N},MATCH(B{r},汇率币种头,0))'
            ws[f'H{r}'] = f'=IF({A}="","",IF(B{r}=本位币,1,IFERROR(IF(N({rt})=0,"",{rt}),"")))'
            ws[f'I{r}'] = f'=IF(OR({A}="",H{r}=""),"",IFERROR(ROUND(G{r}/H{r},2),""))'
        else:
            rt = f'INDEX(记账汇率区,{REP_N},MATCH(B{r},汇率币种头,0))'
            ws[f'H{r}'] = f'=IF({A}="","",IF(B{r}=本位币,1,IFERROR(IF(N({rt})=0,"",{rt}),"")))'
            ws[f'I{r}'] = f'=IF(OR({A}="",H{r}=""),"",ROUND(G{r}*H{r},2))'
        ws[f'J{r}'] = f'=IF({A}="","",IFERROR(INDEX(余额期末,MATCH(C{r},余额编码,0)),0))'
        # opt：对应科目没填时 J 取到的是空（原来 K＝空−数 报 #VALUE!），现在 K 也空着（数据校验照样报这个账户没对上科目）
        ws[f'K{r}'] = f'=IF(OR({A}="",I{r}="",J{r}=""),"",ROUND(J{r}-I{r},2))'
        ws[f'L{r}'] = f'=IF({A}="","",{side(K_IN, *y)})'
        ws[f'M{r}'] = f'=IF({A}="","",{side(K_OUT, *y)})'
    style_rows(ws, R0, R1, cols, auto=tuple(cols),
               fmts={**{c: MONEY for c in 'DEFGIJKLM'}, 'H': RATE},
               aligns={'A': AL, **{c: AR_ for c in 'DEFGHIJKLM'}})
    ws.conditional_formatting.add(f'K{R0}:K{R1}', FormulaRule(formula=[f'AND(ISNUMBER(K{R0}),K{R0}<>0)'], fill=FILL_WARN,
                                                               font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    r = R1 + 1
    put(ws, f'A{r}', '合计（折苏姆）', F_TXTB, FILL_TOT, align=AC)
    for c in 'BCDEFGHLM':
        put(ws, f'{c}{r}', None, F_TXTB, FILL_TOT)
    for c in 'IJK':
        put(ws, f'{c}{r}', f'=SUM({c}{R0}:{c}{R1})', F_TXTB, FILL_TOT, MONEY, AR_)
    put(ws, f'A{r + 2}', '原币的收入、支出、余额只能同一个账户比；跨币种只看「折苏姆」几列。', F_NOTE, border=False)
    ws.freeze_panes = f'B{R0}'
    ctx['acct_rows'] = (R0, R1)
    return ws


# ─────────────────────────── 往来余额表 ───────────────────────────
ARAP_ACCS = [('C', '1122', '应收账款', 1), ('D', '1123', '预付账款', 1), ('E', '1221', '其他应收款\n不含待对冲', 1),
             ('F', '2202', '应付账款', -1), ('G', '2203', '预收账款', -1), ('H', '2241', '其他应付款', -1),
             ('I', '3001', '实收资本', -1)]


def build_arap(wb, ctx):
    ws = wb.create_sheet(SH_ARAP)
    widths(ws, {'A': 20, 'B': 8, **{c: 17 for c, *_ in ARAP_ACCS}, 'J': 17, 'K': 17, 'L': 15, 'M': 13, 'N': 17, 'O': 17, 'P': 16})
    title(ws, '往来余额表（按往来单位 · 报表月末）', 'P', C_LED,
          '💡 客户、供应商、股东、借支的员工，各挂了多少：借方类（应收、预付、其他应收）正数＝人家欠我们；贷方类（应付、预收、其他应付、实收资本）'
          '正数＝我们欠人家/人家投进来的。取【记账分录】里「往来单位」这一栏，所以流水、采购、销售录的时候往来单位要写。'
          '最下面「没写往来单位的」是这些科目里没挂到人的部分。J、K 两列是流水里跟这个单位的本年收付款（折苏姆）。'
          'L～P：人民币、美元进货还欠供应商多少（原币），按月末汇率该是多少苏姆，跟账上差多少——P 列不是 0 就是汇率差，'
          '做张手工凭证转汇兑损益（差额为正：借 应付账款、贷 财务费用-汇兑损益；往来单位写这个供应商）。')
    header(ws, 4, [('A', '往来单位'), ('B', '类型')] + [(c, f'{nm}\n({code})') for c, code, nm, _ in ARAP_ACCS]
           + [('J', '本年收款\n(流水)'), ('K', '本年付款\n(流水)'), ('L', '外币应付\n人民币(原币)'), ('M', '外币应付\n美元(原币)'),
              ('N', '按月末汇率\n应有苏姆'), ('O', '账上外币货款\n苏姆'), ('P', '汇率差\n(应转汇兑损益)')], C_LED, height=40)
    cp = lambda col: rng(SH_BASE, col, CP_R0, CP_R1)
    R0, R1 = 5, 5 + CP_R1 - CP_R0
    end = f'"<="&{REP_M1}'
    for i in range(CP_R1 - CP_R0 + 1):
        r = R0 + i
        k = i + 1
        A = f'$A{r}'
        ws[f'A{r}'] = f'=IF({at(cp(CP_NAME), k)}="","",{at(cp(CP_NAME), k)}&"")'
        ws[f'B{r}'] = f'=IF({A}="","",{at(cp(CP_TYPE), k)}&"")'
        # opt：原来每格 2～6 个 SUMIFS(…,分录往来,单位,分录科目,"1122*",分录日期,"<="&月末)（通配符、1.8 万行，300 个单位都算）。
        # 现在：Q 列＝这个单位在记账分录里有没有（没有就全是 0，不算）；有的按记账分录 W 列「分类键」（1122→11 … 3001→17，
        # 1221 已经去掉 122104、122105）精确汇总到报表月：每格 2 个 SUMIFS、条件都是数字
        ws[f'Q{r}'] = f'=IF({A}="",0,IF(COUNTIF(分录往来,{A})>0,1,0))'
        for c, code, _, sgn in ARAP_ACCS:
            base = CLS_ARAP[code] * 100000
            k = f'分录分类键,">="&{base},分录分类键,"<="&({base}+报表月序*10)'
            v = f'(SUMIFS(分录借方,分录往来,{A},{k})-SUMIFS(分录贷方,分录往来,{A},{k}))'
            ws[f'{c}{r}'] = f'=IF({A}="","",IF($Q{r}=0,0,{v}))' if sgn > 0 else f'=IF({A}="","",IF($Q{r}=0,0,-{v}))'
        y = f'{cash_u(K_DATE)},">="&{REP_Y0},{cash_u(K_DATE)},"<="&{REP_M1}'
        ws[f'J{r}'] = f'=IF({A}="","",SUMIFS({cash_u(K_UZS)},{cash_u(K_CP)},{A},{cash_u(K_UZS)},">0",{y}))'
        ws[f'K{r}'] = f'=IF({A}="","",-SUMIFS({cash_u(K_UZS)},{cash_u(K_CP)},{A},{cash_u(K_UZS)},"<0",{y}))'
        # 外币应付：进货的外币货款 − 从同币种账户付给他、冲应付账款的钱（原币），以及对应的苏姆
        bd = f'{mod_u(SH_BUY, B_DATE)},{end}'
        cd = f'{cash_u(K_DATE)},{end}'
        fx_orig = lambda cur: (f'SUMIFS({mod_u(SH_BUY, B_AMT0)},{mod_u(SH_BUY, B_SUP)},{A},{mod_u(SH_BUY, B_CUR)},"{cur}",{mod_u(SH_BUY, B_VALID)},1,{bd})'
                               f'+SUMIFS({cash_u(K_NET)},{cash_u(K_CP)},{A},{cash_u(K_CUR)},"{cur}",{cash_u(K_CCODE)},应付科目,{cash_u(K_VALID)},1,{cd})')
        fx_uzs = lambda cur: (f'SUMIFS({mod_u(SH_BUY, B_AMT)},{mod_u(SH_BUY, B_SUP)},{A},{mod_u(SH_BUY, B_CUR)},"{cur}",{mod_u(SH_BUY, B_VALID)},1,{bd})'
                              f'+SUMIFS({cash_u(K_UZS)},{cash_u(K_CP)},{A},{cash_u(K_CUR)},"{cur}",{cash_u(K_CCODE)},应付科目,{cash_u(K_VALID)},1,{cd})')
        rate = lambda cur: f'N(INDEX(记账汇率区,{REP_N},MATCH("{cur}",汇率币种头,0)))'
        ws[f'L{r}'] = f'=IF({A}="","",ROUND({fx_orig("CNY")},2))'
        ws[f'M{r}'] = f'=IF({A}="","",ROUND({fx_orig("USD")},2))'
        ws[f'N{r}'] = f'=IF({A}="","",IFERROR(ROUND(L{r}*{rate("CNY")}+M{r}*{rate("USD")},2),""))'
        # 手工凭证里记给这个供应商的应付账款调整（比如已经做过的汇兑损益凭证）也算进账上
        man = (f'SUMIFS({mod_u(SH_MAN, MN_CR)},{mod_u(SH_MAN, MN_CP)},{A},{mod_u(SH_MAN, MN_CODE2)},应付科目,{mod_u(SH_MAN, MN_OK)},1,{mod_u(SH_MAN, MN_VDATE)},{end})'
               f'-SUMIFS({mod_u(SH_MAN, MN_DR)},{mod_u(SH_MAN, MN_CP)},{A},{mod_u(SH_MAN, MN_CODE2)},应付科目,{mod_u(SH_MAN, MN_OK)},1,{mod_u(SH_MAN, MN_VDATE)},{end})')
        ws[f'O{r}'] = f'=IF({A}="","",ROUND({fx_uzs("CNY")}+{fx_uzs("USD")}+{man},2))'
        ws[f'P{r}'] = f'=IF(OR({A}="",N{r}=""),"",IF(AND(L{r}=0,M{r}=0,O{r}=0),"",ROUND(O{r}-N{r},2)))'
    cols = list('ABCDEFGHIJKLMNOP')
    style_rows(ws, R0, R1, cols, auto=tuple(cols), fmts={c: MONEY for c in 'CDEFGHIJKLMNOP'},
               aligns={'A': AL, **{c: AR_ for c in 'CDEFGHIJKLMNOP'}})
    ws.conditional_formatting.add(f'P{R0}:P{R1}', FormulaRule(formula=[f'AND(ISNUMBER(P{R0}),ABS(P{R0})>=1)'], fill=FILL_WARN,
                                                               font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    r = R1 + 1
    put(ws, f'A{r}', '没写往来单位的', F_TXTB, FILL_SUBH, align=AC)
    put(ws, f'B{r}', None, F_TXTB, FILL_SUBH)
    for c, code, _, sgn in ARAP_ACCS:
        base = CLS_ARAP[code] * 100000
        k = f'分录分类键,">="&{base},分录分类键,"<="&({base}+报表月序*10)'
        tot = f'(SUMIFS(分录借方,{k})-SUMIFS(分录贷方,{k}))'           # opt：同上，不按单位
        tot = tot if sgn > 0 else f'-{tot}'
        put(ws, f'{c}{r}', f'={tot}-SUM({c}{R0}:{c}{R1})', F_TXTB, FILL_SUBH, MONEY, AR_)
    for c in 'JKLMNOP':
        put(ws, f'{c}{r}', None, F_TXTB, FILL_SUBH)
    r2 = r + 1
    put(ws, f'A{r2}', '合计（＝科目余额表；其他应收款不含 122104、122105）', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'B{r2}', None, F_TXTB, FILL_TOT)
    for c in 'CDEFGHIJKLMNOP':
        put(ws, f'{c}{r2}', f'=SUM({c}{R0}:{c}{r})', F_TXTB, FILL_TOT, MONEY, AR_)
    put(ws, 'Q4', '记账分录\n里有', F_NOTE, FILL_AUTO, align=ACW)
    for rr in range(R0, R1 + 1):
        ws[f'Q{rr}'].font = F_NOTE
    hide(ws, ['Q'])
    ws.freeze_panes = f'B{R0}'
    ws.auto_filter.ref = f'A4:P{R1}'
    ctx['arap_fx'] = f"{q(SH_ARAP)}!$P${r2}"
    return ws


# ─────────────────────────── 存货收发存 ───────────────────────────
def build_inv(wb, ctx):
    ws = wb.create_sheet(SH_INV)
    cols = 'A B C D E F G H I J K L M N'.split()
    widths(ws, dict(zip(cols, [16, 9, 6, 11, 10, 12, 16, 12, 16, 12, 16, 12, 16, 13])))
    title(ws, '存货收发存（材料 · 按【首页】报表月 · 全月一次加权平均）', 'N', C_LED,
          '💡 只列「库存管理＝是」的物料（炸药、油料、配件、五金、劳保…）。入库取【采购入库】，出库取【领用出库】，'
          '单价＝（月初金额＋本月入库金额）÷（月初数量＋本月入库数量）。原矿的产销存在【原矿产销存】。'
          '月末数量变成负数（红字）＝领的比进的多，先查采购入库有没有漏录。')
    _kpi_band(ws, 'N')
    R0, R1 = 5, 5 + INV_R1 - INV_R0
    _kpi(ws, 'A3', '期间：', 'B3', f'={REP_M0}', MONTH)
    ws.merge_cells('B3:C3')
    _kpi(ws, 'D3', '月末金额合计：', 'F3', f'=SUM(M{R0}:M{R1})', MONEY)
    ws.merge_cells('D3:E3')
    tb = '+'.join(f'N(IFERROR(INDEX(余额期末,MATCH("{c}",余额编码,0)),0))' for c in ('1403', '1411'))
    _kpi(ws, 'H3', '科目余额表 原材料＋周转材料：', 'J3', f'={tb}', MONEY)
    ws.merge_cells('H3:I3')
    _kpi(ws, 'K3', '差额：', 'L3', f'=ROUND(J3-F3,2)', MONEY)
    header(ws, 4, [('A', '物料'), ('B', '规格'), ('C', '单位'), ('D', '类别'), ('E', '存货科目'), ('F', '月初数量'),
                   ('G', '月初金额'), ('H', '本月入库\n数量'), ('I', '本月入库\n金额'), ('J', '本月出库\n数量'),
                   ('K', '本月出库\n金额'), ('L', '月末数量'), ('M', '月末金额'), ('N', '加权单价')], C_LED, height=36)
    mat = lambda col: rng(SH_BASE, col, MAT_R0, MAT_R1)
    n = REP_N
    ok = f'AND({n}>=1,{n}<={N_MONTHS})'
    for i in range(INV_R1 - INV_R0 + 1):
        r = R0 + i
        k = i + 1
        A = f'$A{r}'
        live = f'{at(mat(M_STOCK), k)}="是"'
        ws[f'A{r}'] = f'=IF(AND({at(mat(M_NAME), k)}<>"",{live}),{at(mat(M_NAME), k)}&"","")'
        ws[f'B{r}'] = f'=IF({A}="","",{at(mat(M_SPEC), k)}&"")'
        ws[f'C{r}'] = f'=IF({A}="","",{at(mat(M_UNIT), k)}&"")'
        ws[f'D{r}'] = f'=IF({A}="","",{at(mat(M_CAT), k)}&"")'
        ws[f'E{r}'] = f'=IF({A}="","",{at(mat(M_ACC), k)}&"")'
        g = lambda f, nn=n: inv_field(k, nn, f)
        ws[f'F{r}'] = f'=IF(OR({A}="",NOT({ok})),"",IF({n}=1,N({at(mat(M_Q0), k)}),{g("期末数量", f"{n}-1")}))'
        ws[f'G{r}'] = f'=IF(OR({A}="",NOT({ok})),"",IF({n}=1,N({at(mat(M_A0), k)}),{g("期末金额", f"{n}-1")}))'
        ws[f'H{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("入库数量")})'
        ws[f'I{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("入库金额")})'
        ws[f'J{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("出库数量")})'
        ws[f'L{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("期末数量")})'
        ws[f'M{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("期末金额")})'
        ws[f'K{r}'] = f'=IF(OR({A}="",NOT({ok})),"",ROUND(G{r}+I{r}-M{r},2))'
        ws[f'N{r}'] = f'=IF(OR({A}="",NOT({ok})),"",{g("加权单价")})'
    style_rows(ws, R0, R1, cols, auto=tuple(cols),
               fmts={'F': QTY2, 'G': MONEY, 'H': QTY2, 'I': MONEY, 'J': QTY2, 'K': MONEY, 'L': QTY2, 'M': MONEY, 'N': PRICE},
               aligns={'A': AL, **{c: AR_ for c in 'FGHIJKLMN'}})
    ws.conditional_formatting.add(f'L{R0}:L{R1}', FormulaRule(formula=[f'AND(ISNUMBER(L{R0}),L{R0}<0)'],
                                                               font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    r = R1 + 1
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'BCDEFHJLN':
        put(ws, f'{c}{r}', None, F_TXTB, FILL_TOT)
    for c in 'GIKM':
        put(ws, f'{c}{r}', f'=SUM({c}{R0}:{c}{R1})', F_TXTB, FILL_TOT, MONEY, AR_)
    ws.freeze_panes = f'B{R0}'
    ws.auto_filter.ref = f'A4:N{R1}'
    ctx['inv_rows'] = (R0, R1)
    return ws
