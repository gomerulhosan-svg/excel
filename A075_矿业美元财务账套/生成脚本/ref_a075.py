# -*- coding: utf-8 -*-
"""从用户上传的 A075 原件（../参考/A075_用户上传原件-2026.10.14.xlsx）里读「用户的东西」。

分工（写给以后改这份脚本的人）：
- 公式、下拉、条件格式、名称、容量、表的结构 —— 全部由生成脚本（common.py + s_*.py + p1008.py）按代码写出来；
- 用户填的 / 用户自己排的 —— 从原件按位置读：参数、资金账户、项目、往来单位、物料、记账规则各张对照表、科目表、
  月度汇率、流水（现在 0 行）、明细账默认科目；【首页】【使用说明】两张是用户自己重排过版面的，整张照搬（见 copy_sheet）；
  现金流水总表 T 列的旧批注也照搬（A062 鑫盛那几笔的说明，A075 用不上，见汇报「挡住使用的问题」第 2 条）。
原件只读，不改。读之前不执行原件里的任何东西（openpyxl 只解析 XML）。
10/15：read() 最后调 apply_1015()，按用户的手工登记表改账户对应科目和期初、收支类别（方案甲＋二级明细对照）、物料、科目表，
并把资金台帐 999 笔流水、货柜单号放进来（数据在 导入/*.json，读法在 导入/u4.py；第三方模板的东西一个字不进仓库）。
"""
import datetime as dt
import os
from copy import copy

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from common import *

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
REF = os.path.join(ROOT, '参考', 'A075_用户上传原件-2026.10.14.xlsx')


def load(path=REF):
    return openpyxl.load_workbook(path)


def _v(ws, ref):
    v = ws[ref].value
    if isinstance(v, str) and v.startswith('='):
        return None            # 输入区里不该有公式；有的话不当数据
    return v


def _rows(ws, cols, r0, r1):
    return [tuple(_v(ws, f'{c}{r}') for c in cols) for r in range(r0, r1 + 1)]


class Fixes:
    """顶替 A062 的 fixes.py：同名属性，内容全部来自原件"""


def read(ref):
    b, ru, ra, ac = ref[SH_BASE], ref[SH_RULE], ref[SH_RATE], ref[SH_ACC]
    params = {k: _v(b, f'B{P_ROW[k]}') for k, _, _ in PARAMS}
    fx = Fixes()
    fx.COMPANY, fx.GOLIVE, fx.VAT = params['company'], params['golive'], params['vat']
    accounts = []
    for r in range(ACC_R0, ACC_R1 + 1):
        accounts.append(dict(name=_v(b, f'{A_NAME}{r}'), bank=_v(b, f'{A_BANK}{r}'), cur=_v(b, f'{A_CUR}{r}'),
                             open=_v(b, f'{A_OPEN}{r}'), openb=_v(b, f'{A_OPENB}{r}'), code=_v(b, f'{A_CODE}{r}'),
                             note=_v(b, f'{A_NOTE}{r}')))
    fx.ACCOUNT_CODE = {a['name']: a['code'] for a in accounts if a['name']}
    projects = [_v(b, f'{PJ_NAME}{r}') for r in range(PRJ_R0, PRJ_R1 + 1)]
    cps = [dict(name=_v(b, f'{CP_NAME}{r}'), type=_v(b, f'{CP_TYPE}{r}'), cap=_v(b, f'{CP_CAP}{r}'), note=_v(b, f'{CP_NOTE}{r}'))
           for r in range(CP_R0, CP_R1 + 1)]
    fx.MATERIALS = [(_v(b, f'{M_NAME}{r}'), _v(b, f'{M_SPEC}{r}'), _v(b, f'{M_UNIT}{r}'), _v(b, f'{M_CAT}{r}'),
                     _v(b, f'{M_Q0}{r}'), _v(b, f'{M_P0}{r}'), _v(b, f'{M_NOTE}{r}')) for r in range(MAT_R0, MAT_R1 + 1)]
    rules = _rows(ru, [RC_NAME, RC_ATTR, RC_IN, RC_OUT, RC_HIST, RC_CFIN, RC_CFOUT, RC_NOTE], RC_R0, RC_R1)
    fx.MAT_CATS = _rows(ru, [MC_NAME, MC_ACC, MC_STOCK, MC_PROD, MC_NOTE], MC_R0, MC_R1)
    fx.USES = _rows(ru, [US_NAME, US_ACC, US_NOTE], US_R0, US_R1)
    fx.DEPTS = _rows(ru, [DP_NAME, DP_WAGE, DP_DEP, DP_NOTE], DP_R0, DP_R1)
    fx.PERSON_TYPES = _rows(ru, [PT_NAME, PT_ACC], PT_R0, PT_R1)
    fx.FA_CATS = _rows(ru, [FC_NAME, FC_ACC, FC_YEARS, FC_SALV, FC_NOTE], FC_R0, FC_R1)
    fx.CF_ITEMS = _rows(ru, [CFI_NO, CFI_NAME, CFI_SIDE], CFI_R0, CFI_R1)
    chart = []
    for r in range(AC_R0, AC_R1 + 1):
        code = _v(ac, f'{AC_CODE}{r}')
        a = dict(code=None if code is None else str(code), name=_v(ac, f'{AC_NAME}{r}'), cls=_v(ac, f'{AC_CLASS}{r}'),
                 dir=_v(ac, f'{AC_DIR}{r}'), aux=_v(ac, f'{AC_AUX}{r}'), fx=_v(ac, f'{AC_FX}{r}'),
                 cashf=_v(ac, f'{AC_CASHF}{r}'), stat=_v(ac, f'{AC_STAT}{r}'), para=_v(ac, f'{AC_PARA}{r}'),
                 open=_v(ac, f'{AC_OPEN}{r}'),
                 added=bool(ac[f'{AC_NAME}{r}'].font and ac[f'{AC_NAME}{r}'].font.b))
        chart.append(a if code is not None else {})
    rates = []
    for r in range(RATE_R0, RATE_R1 + 1):
        rates.append({c: _v(ra, f'{CL(i + 2)}{r}') for i, c in enumerate(RATE_CURS)})
    # 流水（A075 现在 0 行；有的话照搬录入列）
    cs = ref[SH_CASH]
    ledger = []
    inp = dict(date=K_DATE, account=K_ACC, category=K_CAT, memo=K_MEMO, inc=K_IN, exp=K_OUT, xrate=K_XR,
               project=K_PRJ, cp=K_CP, note=K_NOTE, ovr=K_OVR, merge=K_MERGE, att=K_ATT)
    for r in range(CASH_R0, CASH_R1 + 1):
        d = {k: _v(cs, f'{c}{r}') for k, c in inp.items()}
        if any(v is not None for v in d.values()):
            d['row'] = r
            ledger.append(d)
    start = params['start']
    if isinstance(start, (int, float)):        # WPS 存的日期格式号（57）openpyxl 不认，读出来是序列号
        from openpyxl.utils.datetime import from_excel
        start = from_excel(start)
    months = [dt.datetime(start.year + (start.month - 1 + i) // 12, (start.month - 1 + i) % 12 + 1, 1) for i in range(N_MONTHS)]
    out = dict(params=params, fx=fx, accounts=accounts, projects=projects, counterparties=cps, rules=rules, chart=chart,
               rates=rates, ledger=ledger, months=months)
    apply_1015(out)
    return out


# ───────────────────────── 10/15：按用户的手工登记表调整 ＋ 导入资金台帐（plan.md 第二节） ─────────────────────────
# 用户 10/15 上传了三本手工登记表和一本资金台帐（u4）。这里把原件读出来的东西按她的叫法、她的物资、她的账户改一遍，
# 再把资金台帐 999 笔流水放进 ctx['ledger']（数据在 导入/，读法在 导入/u4.py）。改了什么都在下面几张表里，改就改这里。

# 资金账户（【基础资料】E5:J7）：对应科目照她老账（美元＝库存现金、FBC＝银行存款-津币、人民币＝其他货币资金），
# 期初原币、期初折美元照资金台帐；【科目表】M 列期初按同一个数填（两处一样，资金账户余额表 K 列才是 0）
ACC_CODE_1015 = {'人民币': '101202', '美元': '100105', 'FBC': '100210'}
ACC_NOTE_1015 = {'人民币': '国内的人民币（老账记「其他货币资金-人民币」）', '美元': '美元现金（老账记「库存现金-美金」）',
                 'FBC': 'FBC 银行津币户（老账记「银行存款-津币」）'}
PLUG_CODE = '310499'               # 期初只有资金账户：差额先放这里，科目余额表、资产负债表才平（数据校验另有一项提醒）
PARAMS_1015 = {'vat': 0, 'vatin': '否'}       # 卖黄金不交销项税（底稿帐里卖金没记销项税）
BUYERS = [('私人', '客户', '黄金买方'), ('银行', '客户', '黄金买方'), ('银行代理', '客户', '黄金买方')]   # 【销售结算】D 列「买方」下拉

# 【记账规则】收支类别：方案甲，沿用她自己的叫法（资金台帐 u4 里用的类别＝她老账的一级科目）
# (名称, 属性, 收到钱时, 付出钱时, 启用日前, 现金流量收, 现金流量付, 说明)
RULE_FIX = {   # 原有的类别改一下
    '销售回款': dict(hist='500103', note='买方付黄金款：冲应收账款-黄金销售（先收钱后结算的，挂成预收）。业务模块启用日以前的直接记主营业务收入-黄金销售'),
}
_MAT = '材料款：业务模块启用日以前直接算成本（生产耗用）；启用日起冲应付账款（货在【采购入库】登记）'
RULES_ADD = [
    ('主营业务收入', '收入', '112203', '112203', '500103', 1, 1, '卖黄金收到的钱。启用日以前直接记主营业务收入-黄金销售；启用日起冲应收账款（卖金在【销售结算】登记）'),
    ('原材料', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('柴油', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('氰化钠', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('炸药', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('引线', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('盐酸', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('硝酸', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('烧碱', '支出', '220202', '220202', '540102', 3, 3, _MAT),
    ('材料差异', '支出', '224105', '224105', '540102', 3, 3, '清关费、运费、空运费：启用日以前直接算成本；启用日起冲应付进口税费及运杂'),
    ('生活支出', '支出', '56021001', '56021001', None, 6, 6, '食堂、矿泉水、饭费'),
    ('应交税费', '支出', '560220', '560220', None, 5, 5, '交给政府的税费；按「二级」分到明细（PAYE、NSSA、土地税…见右边「二级明细对照」）'),
    ('待摊费用', '支出', '1801', '1801', None, 6, 6, '堆淋、氰化厂的征地、搬迁、赔偿等，记长期待摊费用'),
    ('长期待摊费用', '支出', '1801', '1801', None, 12, 12, '围栏等'),
    ('无形资产', '支出', '170101', '170101', None, 12, 12, '矿权记采矿权；征地款「二级」写「土地」记土地使用权'),
    ('销售费用', '支出', '560199', '560199', None, 6, 6, '卖金时银行扣的钱等；银行手续费「二级」写「银行扣费」记财务费用-手续费'),
]

# 「二级明细对照」（【记账规则】最右边）：类别＋二级（流水 R 列）→ 记到哪个科目。只给跟启用日无关的类别配；
# 材料类、工资（应付职工薪酬-工资）、黄金收入启用日前后记的科目不一样，不配（照收支类别那一块走）
SUBS = [   # (类别, 二级, 科目, 说明)
    ('管理费用', '其他', '560299', None),
    ('管理费用', '安保费', '560299', '要单设科目再改'),
    ('管理费用', '公关费', '560205', '业务招待费'),
    ('管理费用', '签证差旅费', '560212', '机票及签证费'),
    ('管理费用', '小车及摩托车费', '56020604', '车辆费用-其他'),
    ('管理费用', '医药费', '56021003', '福利费-其他'),
    ('管理费用', '伙食费', '56021001', None),
    ('管理费用', '餐费', '56021001', None),
    ('管理费用', '维修费', '560214', None),
    ('管理费用', '修理费', '560214', None),
    ('管理费用', '网络邮电费', '56020802', None),
    ('管理费用', '电费', '56020402', None),
    ('管理费用', '办公费', '560201', None),
    ('应付职工薪酬', '职工福利', '56021003', '年终奖、过节费、礼金'),
    ('应交税费', 'PAYE', '540399', '员工个税。启用日以后用了【工资计提】的，个税改选类别「代扣个人所得税」'),
    ('应交税费', 'NSSA', '56021301', '社保（单位负担）'),
    ('应交税费', '土地税', '540319', None),
    ('应交税费', '打井抽水许可', '540317', '水资源使用费'),
    ('应交税费', 'SUB FEES', '560220', None),
    ('应交税费', 'ZIMDEF', '560220', None),
    ('应交税费', 'NEC', '560220', None),
    ('应交税费', 'STANDARDS', '560220', None),
    ('应交税费', '许可证', '560220', None),
    ('应交税费', '环保费', '560220', None),
    ('应交税费', '劳工工会费', '560220', None),
    ('固定资产', '宽体车', '160101', '矿用车辆'),
    ('固定资产', '洒水车', '160107', '运输及乘用车辆'),
    ('固定资产', '货箱车', '160107', None),
    ('固定资产', '5吨轻卡车', '160107', None),
    ('固定资产', '电解设备', '160106', '其他机械设备'),
    ('无形资产', '矿权', '170101', '采矿权'),
    ('无形资产', '土地', '170102', '土地使用权（征地款）'),
    ('销售费用', '银行扣款', '560199', '卖金时银行扣的约 5%，是不是权利金待确认（是的话改 540311）'),
    ('销售费用', '银行扣费', '560302', '财务费用-手续费'),
    ('销售费用', '压缩机', '560216', '低值易耗品'),
    ('其他应收款', '三一', '122106', '保证金（光伏）'),
    ('待摊费用', '堆淋-其他', '1801', None),
    ('长期待摊费用', '围栏', '1801', None),
]

# 物料类别、领用用途（【记账规则】）：加她的
MAT_CATS_ADD = [   # (类别, 入账科目, 库存管理, 采矿生产领用时计入, 说明)
    ('化工药剂', '140306', '是', '40010103', '氰化钠、盐酸、硝酸、硫酸、片碱、活性炭、石灰'),
    ('建材', '40010109', '否', '40010109', '水泥、红砖、石子：到货就用（修路修桥、码沟），买来直接算生产成本，不进库'),
]
USES_ADD = [   # (用途, 计入科目, 说明)
    ('堆淋', '按物料类别', '堆淋场用的（堆号写在「项目」）'),
    ('氰化厂', '按物料类别', '氰化厂、出碳用的'),
    ('修路修桥', '40010109', '修路、修桥、码沟'),
    ('借出/抵消', '122103', '借给别的矿山或抵消的：记其他应收款，不算成本'),
]
# 物料档案（【基础资料】X:AH）：换成她的物资（手工出入库台账、进货台账里的），A062 的通用材料和设备留着
# (名称, 规格, 单位, 类别, 期初数量, 期初单价, 备注)
MATERIALS_1015 = [
    ('炸药', None, '箱', '火工品', None, None, '1 箱＝25 公斤'),
    ('雷管', None, '发', '火工品', None, None, None),
    ('导爆管', None, '发', '火工品', None, None, '15 米、8 米'),
    ('导爆索', None, '卷', '火工品', None, None, '引线，250 米/卷'),
    ('引火线', None, '卷', '火工品', None, None, None),
    ('氰化钠', None, '箱', '化工药剂', None, None, '1 箱约 1.1 吨'),
    ('盐酸', None, '桶', '化工药剂', None, None, None),
    ('硝酸', None, '桶', '化工药剂', None, None, None),
    ('硫酸', None, '桶', '化工药剂', None, None, None),
    ('片碱', None, '包', '化工药剂', None, None, '烧碱，1 包＝25 公斤'),
    ('活性炭', None, '包', '化工药剂', None, None, None),
    ('石灰', None, '吨', '化工药剂', None, None, None),
    ('柴油', None, '升', '油料', None, None, None),
    ('水泥', None, '包', '建材', None, None, '325 / 425；600 包＝30 吨'),
    ('红砖', None, '块', '建材', None, None, None),
    ('石子', None, '方', '建材', None, None, None),
    ('彩条布', None, '卷', '其他材料', None, None, None),
]
MATERIALS_REPLACED = {'乳化炸药', '雷管', '导爆索', '柴油'}     # A062 的这几样换成上面她的（单位不一样）

# 【科目表】：改名 / 新增 / 删掉 A062 的苏姆科目；「外币核算」列的 SOM（苏姆）改成美元口径
CHART_DROP = {'100101', '100102', '100207', '100208'}          # 苏姆现金、苏姆工资卡、NBU 苏姆结算户、公务卡苏姆户（A062 乌兹别克的）
CHART_RENAME = {
    '112203': '黄金销售', '140501': '黄金', '500103': '黄金销售', '540101': '黄金销售成本',
    '540102': '生产耗用（启用日前现金口径）', '170101': '采矿权', '163103': '宿舍区', '163104': '管理区',
}
CHART_ADD = [   # 加在最后（跟用户以前加 122106、100209、100210 一样）
    dict(code='101202', name='人民币', cls='流动资产', dir='借', fx='CNY', cashf='是'),
    dict(code='140306', name='化工药剂', cls='流动资产', dir='借', fx='USD', cashf='否'),
    dict(code='170102', name='土地使用权（征地款）', cls='非流动资产', dir='借', fx='USD', cashf='否'),
    dict(code=PLUG_CODE, name='期初差额待核', cls='所有者权益', dir='贷', fx='USD', cashf='否'),
]
FX_TEXT = {'SOM/RMB/USD': 'USD/CNY/ZIG', 'SOM': 'USD'}


def _u4():
    import importlib.util
    path = os.path.join(HERE, '导入', 'u4.py')
    spec = importlib.util.spec_from_file_location('a075_import_u4', path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def apply_1015(d):
    """read() 读出来的 d 就地改成 10/15 的样子（d 的键见 read）；另加 d['subs']（二级明细对照）、d['containers']（货柜登记）"""
    u4 = _u4()
    p, fx = d['params'], d['fx']
    p.update(PARAMS_1015)
    fx.VAT = p['vat']
    # 资金账户
    accs = {a['name']: a for a in u4.accounts()}
    for a in d['accounts']:
        if a.get('name') in accs:
            x = accs[a['name']]
            assert (a['bank'], a['cur']) == (x['bank'], x['cur']), (a, x)
            a.update(open=x['open'], openb=x['openb'], code=ACC_CODE_1015[a['name']], note=ACC_NOTE_1015[a['name']])
    assert sum(1 for a in d['accounts'] if a.get('code')) == len(accs) == 3
    fx.ACCOUNT_CODE = {a['name']: a['code'] for a in d['accounts'] if a['name']}
    # 往来单位：黄金买方
    names = {c['name'] for c in d['counterparties'] if c.get('name')}
    free = [c for c in d['counterparties'] if not c.get('name')]
    for (nm, tp, note), c in zip([b for b in BUYERS if b[0] not in names], free):
        c.update(name=nm, type=tp, note=note)
    # 收支类别
    rules = [list(x) for x in d['rules']]
    have = {x[0]: x for x in rules if x[0]}
    for nm, fix in RULE_FIX.items():
        have[nm][4] = fix['hist']
        have[nm][7] = fix['note']
    n = sum(1 for x in rules if x[0])
    assert all(x[0] is None for x in rules[n:]), '记账规则收支类别中间有空行'
    for x in RULES_ADD:
        assert x[0] not in have, x[0]
        rules[n] = list(x)
        n += 1
    d['rules'] = [tuple(x) for x in rules]
    d['subs'] = list(SUBS)
    # 物料类别、用途
    def _append(lst, add, width):
        lst = list(lst)
        k = sum(1 for x in lst if x[0])
        assert all(x[0] is None for x in lst[k:]) and k + len(add) <= len(lst)
        lst[k:k + len(add)] = [tuple(x) + (None,) * (width - len(x)) for x in add]
        return lst
    fx.MAT_CATS = _append(fx.MAT_CATS, MAT_CATS_ADD, 5)
    fx.USES = _append(fx.USES, USES_ADD, 3)
    # 物料档案
    old = [m for m in fx.MATERIALS if m[0] and m[0] not in MATERIALS_REPLACED]
    mats = MATERIALS_1015 + old
    fx.MATERIALS = mats + [(None,) * 7] * (len(fx.MATERIALS) - len(mats))
    assert len(fx.MATERIALS) == MAT_R1 - MAT_R0 + 1
    # 科目表
    chart = [a for a in d['chart'] if a and a['code'] not in CHART_DROP]
    assert len([a for a in d['chart'] if a and a['code'] in CHART_DROP]) == len(CHART_DROP)
    codes = {a['code'] for a in chart}
    for a in chart:
        if a['code'] in CHART_RENAME:
            a['name'] = CHART_RENAME[a['code']]
            a['added'] = True
        a['fx'] = FX_TEXT.get(a['fx'], a['fx'])
    for x in CHART_ADD:
        assert x['code'] not in codes, x
        chart.append(dict(x, aux=None, stat='启用', para='否', open=None, added=True))
    opens = {a['code']: a['openb'] for a in d['accounts'] if a.get('code')}
    for a in chart:
        if a['code'] in opens:
            assert a.get('open') in (None, 0), a
            a['open'] = opens[a['code']]
    by = {a['code']: a for a in chart}
    by[PLUG_CODE]['open'] = round(-sum(opens.values()), 2)
    assert len(chart) <= AC_R1 - AC_R0 + 1
    d['chart'] = chart + [{}] * (AC_R1 - AC_R0 + 1 - len(chart))
    # 流水：资金台帐 999 笔（原件里 0 笔）
    assert not d['ledger'], '原件里有流水了，导入要另想办法'
    d['ledger'] = u4.ledger()
    d['containers'] = u4.containers()
    return d


# ───────────────────────── 整张照搬（首页、使用说明：用户自己重排过的版面） ─────────────────────────
def copy_sheet(ref, wb, name, index=None, drop_rows=()):
    """整张照搬；drop_rows：不要的行（下面的行往上挪，合并格、行高跟着挪；有效性、条件格式只能在这些行上面）"""
    src = ref[name]
    if name in wb.sheetnames:
        old = wb[name]
        index = wb.sheetnames.index(name) if index is None else index
        wb.remove(old)
    ws = wb.create_sheet(name, index)
    drop = sorted(drop_rows)
    nr = lambda r: r - sum(1 for x in drop if x < r)
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            if c.row in drop:
                continue
            d = ws.cell(row=nr(c.row), column=c.column, value=c.value)
            if c.has_style:
                d.font, d.fill, d.border = copy(c.font), copy(c.fill), copy(c.border)
                d.alignment, d.protection = copy(c.alignment), copy(c.protection)
                d.number_format = c.number_format
            if c.hyperlink is not None:
                d.hyperlink = copy(c.hyperlink)
                d.hyperlink.ref = d.coordinate
    for mr in src.merged_cells.ranges:
        if any(mr.min_row <= x <= mr.max_row for x in drop):
            assert mr.min_row == mr.max_row, (name, str(mr))
            continue
        ws.merge_cells(start_row=nr(mr.min_row), start_column=mr.min_col, end_row=nr(mr.max_row), end_column=mr.max_col)
    for k, cd in src.column_dimensions.items():
        n = ws.column_dimensions[k]
        n.width, n.hidden = cd.width, cd.hidden
        n.min, n.max = cd.min, cd.max
        if cd.outlineLevel:
            n.outlineLevel = cd.outlineLevel
    for k, rd in src.row_dimensions.items():
        if k in drop:
            continue
        n = ws.row_dimensions[nr(k)]
        if rd.height is not None:
            n.height = rd.height
        n.hidden = rd.hidden
    if drop:
        lo = min(drop)
        assert all(r.max_row < lo for dv in src.data_validations.dataValidation for r in dv.sqref.ranges), name
        assert all(r.max_row < lo for cf in src.conditional_formatting for r in cf.sqref.ranges), name
    for dv in src.data_validations.dataValidation:
        nd = DataValidation(type=dv.type, operator=dv.operator, formula1=dv.formula1, formula2=dv.formula2,
                            allow_blank=dv.allow_blank, showErrorMessage=dv.showErrorMessage,
                            showInputMessage=dv.showInputMessage, errorStyle=dv.errorStyle,
                            errorTitle=dv.errorTitle, error=dv.error, promptTitle=dv.promptTitle, prompt=dv.prompt)
        ws.add_data_validation(nd)
        nd.sqref = dv.sqref
    for rng_ in src.conditional_formatting:
        for rule in rng_.rules:
            ws.conditional_formatting.add(str(rng_.sqref), rule)
    ws.sheet_properties.tabColor = copy(src.sheet_properties.tabColor)
    sv, dv_ = src.sheet_view, ws.sheet_view
    dv_.showGridLines, dv_.zoomScale = sv.showGridLines, sv.zoomScale
    if src.freeze_panes:
        ws.freeze_panes = src.freeze_panes
    ws.page_setup.orientation, ws.page_setup.paperSize = src.page_setup.orientation, src.page_setup.paperSize
    return ws


def copy_comments(ref, wb, name, col):
    """把原件某一列的批注照搬过来（作者、文字、框大小）"""
    from openpyxl.comments import Comment
    src, ws = ref[name], wb[name]
    n = 0
    for row in src.iter_rows(min_col=CI(col), max_col=CI(col)):
        for c in row:
            if c.comment is not None:
                cm = Comment(c.comment.text, c.comment.author)
                cm.width, cm.height = c.comment.width, c.comment.height
                ws[c.coordinate].comment = cm
                n += 1
    return n
