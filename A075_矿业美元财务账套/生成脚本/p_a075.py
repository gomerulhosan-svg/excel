# -*- coding: utf-8 -*-
"""A075 的改动里「不好写成参数」的那些，接在 p1008.py 后面执行。

（能写成参数的已经在 common.py「A075 参数」和各 s_*.py 里：本位币美元、汇率方向、月度汇率 K:R、下拉名称…）
这里做：
  · 【首页】【使用说明】整张从原件照搬——这两张用户（10/05）自己重排过版面，A075 又改了文字，
    生成脚本不再自己排（公式也照搬：首页 10 个 KPI 公式、导航链接）。
  · （现金流水总表 T 列 57 条 A062 鑫盛的旧批注：wps 起不再照搬。）
  · 10/15（按用户的手工登记表调整，plan.md 第二节）：fix_words 把「苏姆」→「美元」、「原矿 / 吨 / 甲方 / 出矿」→「黄金 / 克 / 买方 / 冶炼」
    （标题、提示、表头、校验提示、凭证摘要；表名「产量登记」「原矿产销存」不改，只改字）；TEXTS_1015 黄金销售、冶炼登记、二级/备注等表头；
    往来余额表 M 列外币应付改 ZiG；首页改成报表本用的说法；【使用说明】两份（build_help_report＝报表本 / 一本版，build_help_entry＝录入本）；
    【货柜登记】（s_base.build_container）。
  · fix_texts：首页 / 使用说明里过时的文字（数据导入、五个账户、1.6 万行、L3 选科目、编号）；基础资料 J5:J7 对应科目先填建议值。
  · TEXTS：A075 改写过的标题、提示、表头文字（A062 是「苏姆」口径，A075 改成「美元」「1 美元 = X 外币」）。
    右边注释是 A062 / 生成脚本原来的文字。改字就改这里。
  · 用户在 WPS 里调过的「样子」（从原件比出来的，写死在下面几个表里）：标签颜色（用户 10/05 按设置 / 录入 / 自动 / 报表 / 校验分了五种颜色）、
    几列列宽、几行行高、打印设置（纸张、方向、缩放）、现金流水总表 90% 显示、三张表的筛选范围（用户 10/05 往下多拉了 1～2 行）。
"""
from openpyxl.workbook.defined_name import DefinedName

import ref_a075
from common import *

# ───── 用户调过的样子（从原件比出来的） ─────
TAB_COLORS = {
    '基础资料': 'FFFFC000',
    '记账规则': 'FFFFC000',
    '月度汇率': 'FFFFC000',
    '科目表': 'FFFFC000',
    '现金流水总表': 'FF00B050',
    '采购入库': 'FF00B050',
    '领用出库': 'FF00B050',
    '产量登记': 'FF00B050',
    '销售结算': 'FF00B050',
    '工资计提': 'FF00B050',
    '固定资产': 'FF00B050',
    '手工凭证': 'FF00B050',
    '货柜登记': 'FF00B050',
    '月末结转': 'FF0070C0',
    '记账分录': 'FF0070C0',
    '凭证汇总': 'FF0070C0',
    '记账凭证': 'FF0070C0',
    '科目余额表': 'FF7030A0',
    '明细账': 'FF7030A0',
    '资金账户余额表': 'FF7030A0',
    '往来余额表': 'FF7030A0',
    '存货收发存': 'FF7030A0',
    '原矿产销存': 'FF0070C0',
    '资产负债表': 'FF7030A0',
    '利润表': 'FF7030A0',
    '现金流量表': 'FF7030A0',
    '经营报表': 'FF7030A0',
    '数据校验': 'FF7F7F7F',
}
COL_WIDTHS = {
    '科目表': {'C': (16.11, False)},
    '现金流水总表': {'B': (23.43, False), 'E': (16.21, False), 'F': (8.84, False), 'S': (24.68, False), 'U': (33.73, False)},
    '凭证汇总': {'B': (12.89, False), 'D': (48.42, False), 'F': (21.89, False)},
}
ROW_HEIGHTS = {
    '基础资料': {4: 54.0, 31: 14.1, 32: 14.1, 33: 14.1, 34: 14.1, 35: 14.1, 36: 14.1},
    '现金流水总表': {5: 60.0},
    '月末结转': {5: 48.0},
    '往来余额表': {4: 49.0},
}
PAGE_SETUP = {
    '现金流水总表': {'setup': {'paperSize': '9', 'orientation': 'portrait'}, 'fitToPage': None},
    '产量登记': {'setup': {'paperSize': '9', 'orientation': 'portrait'}, 'fitToPage': None},
    '工资计提': {'setup': {'paperSize': '9', 'orientation': 'portrait'}, 'fitToPage': None},
    '资金账户余额表': {'setup': {'paperSize': '9', 'orientation': 'portrait'}, 'fitToPage': None},
    '往来余额表': {'setup': {'paperSize': '9', 'scale': '51', 'fitToHeight': '0', 'orientation': 'landscape'}, 'fitToPage': '1'},
    '资产负债表': {'setup': {'paperSize': '1', 'scale': '76', 'fitToWidth': '0', 'orientation': 'landscape'}, 'fitToPage': None},
    '利润表': {'setup': {'paperSize': '1', 'scale': '93', 'orientation': 'portrait'}, 'fitToPage': '1'},
    '现金流量表': {'setup': {'paperSize': '1', 'scale': '88', 'orientation': 'portrait'}, 'fitToPage': '1'},
    '经营报表': {'setup': {'paperSize': '9', 'scale': '49', 'orientation': 'landscape'}, 'fitToPage': '1'},
}
ZOOM = {
    '现金流水总表': '90',
}
FILTERS = {
    '科目余额表': 'A6:M457',
    '往来余额表': 'A4:P306',
    '存货收发存': 'A4:N205',
}


# ───── A075 改写过的文字（从原件逐格比出来的，58 格） ─────
TEXTS = {
    ('科目表', 'A2'): '💡 编码按 4-2-2-2 分级。末级科目“期初余额”在建账起始日登记（借正贷负），包括已对接的三个资金账户；其他期初资产、负债、权益需按完整试算平衡表补齐。灰色列自动计算，不要删除预留公式。',    # 原来：'💡 编码按 4-2-2-2 分级。要加明细科目：往下面空行填（编码写成文本，比如 160110），不要插行删行。「期初
    ('现金流水总表', 'A1'): '现金流水总表 · 流水录入（美元本位币，账户原币按汇率折美元）',    # 原来：'现金流水总表 · 流水录入（五个账户全量合并，本位币苏姆）'
    ('现金流水总表', 'A2'): '💡 每笔一行：选账户→自动带出币种；USD汇率为1，CNY/ZIG等填“1美元兑换多少原币”；原币净额÷汇率＝折美元。有实际成交汇率就在 J 列填同方向报价。填日期、类别和金额后自动出凭证，全年预留行与公式已保留。不要在中间插入或删除整行。',    # 原来：'💡 一笔一行照原来记：选账户名称 → 银行、币种、汇率自动带出来；收入和支出分两列，只填一个；有实际成交价（购汇、结汇
    ('现金流水总表', 'E3'): '资金账户账面美元合计：',    # 原来：'五个账户账面苏姆合计：'
    ('现金流水总表', 'J4'): '汇率 · 原币折美元（自动）',    # 原来：'汇率 · 折苏姆（自动）'
    ('现金流水总表', 'J5'): '实际成交汇率\n(1美元=X外币)',    # 原来：'实际成交汇率\n(选填)'
    ('现金流水总表', 'K5'): '应用汇率\n(1美元=X外币)',    # 原来：'应用汇率'
    ('现金流水总表', 'M5'): '折美元',    # 原来：'折苏姆'
    ('现金流水总表', 'O5'): '账面美元余额',    # 原来：'账面苏姆余额'
    ('月末结转', 'A2'): '💡 不用填，每个月自动算，凭证自动进【记账分录】（日期是月末最后一天）：① 折旧：【固定资产】卡片按部门汇总；② 调汇：人民币、美元账户按当月汇率重新折美元，差额记财务费用-汇兑损益；③ 生产成本结转：有出矿的月份，把生产成本（炸药、油料、配件、人工、折旧…）全部转成库存商品-原矿；④ 结转销售成本：卖了多少吨按每吨成本转主营业务成本；⑤ 外币兑换差额：【基础资料】参数选「是」才转汇兑损益；⑥ 结转损益：每月最后一天把所有损益科目（收入、成本、税金、费用、营业外收支、所得税）的本月发生额转进 3103 本年利润，损益科目月末余额为 0（J 列＝本月转进本年利润的利润，负数是亏损）；⑦ 年末结转：12 月 31 日把本年利润全部转进 310415 利润分配-未分配利润（K 列）。利润表、经营报表照常按本月 / 本年发生额出数（不含结转凭证），数不变。',    # 原来：'💡 不用填，每个月自动算，凭证自动进【记账分录】（日期是月末最后一天）：① 折旧：【固定资产】卡片按部门汇总；② 调汇
    ('资金账户余额表', 'A1'): '资金账户余额表（原币 · 折美元 · 跟会计科目核对）',    # 原来：'资金账户余额表（原币 · 折苏姆 · 跟账上核对）'
    ('资金账户余额表', 'A2'): '💡 资金账户按月汇总，外币按日记账相同口径（1 USD = X 外币）折美元；若月末账簿余额与报表不一致，检查汇率、调汇、流水核对及其他期初科目。',    # 原来：'💡 按【首页】报表月：每个账户月初、本月收支、月末原币余额（取【现金流水总表】），按当月记账汇率折苏姆，跟科目余额表里
    ('资金账户余额表', 'H4'): '月末汇率\n(1USD=X原币)',    # 原来：'月末\n记账汇率'
    ('资金账户余额表', 'I4'): '月末折美元\n(原币÷汇率)',    # 原来：'月末折苏姆\n(按汇率)'
    ('资金账户余额表', 'J4'): '账面美元\n(科目余额表)',    # 原来：'账上苏姆\n(科目余额表)'
    ('资金账户余额表', 'A25'): '合计（折美元）',    # 原来：'合计（折苏姆）'
    ('资金账户余额表', 'A27'): '原币的收入、支出、余额只能同一个账户比；跨币种只看「折美元」几列。',    # 原来：'原币的收入、支出、余额只能同一个账户比；跨币种只看「折苏姆」几列。'
    ('往来余额表', 'A2'): '💡 客户、供应商、股东、借支的员工，各挂了多少：借方类（应收、预付、其他应收）正数＝人家欠我们；贷方类（应付、预收、其他应付、实收资本）正数＝我们欠人家/人家投进来的。取【记账分录】里「往来单位」这一栏，所以流水、采购、销售录的时候往来单位要写。最下面「没写往来单位的」是这些科目里没挂到人的部分。J、K 两列是流水里跟这个单位的本年收付款（折美元）。L～P：人民币、美元进货还欠供应商多少（原币），按月末汇率该是多少苏姆，跟账上差多少——P 列不是 0 就是汇率差，做张手工凭证转汇兑损益（差额为正：借 应付账款、贷 财务费用-汇兑损益；往来单位写这个供应商）。',    # 原来：'💡 客户、供应商、股东、借支的员工，各挂了多少：借方类（应收、预付、其他应收）正数＝人家欠我们；贷方类（应付、预收、其
    ('原矿产销存', 'A1'): '原矿产销存（吨 · 美元）· 全自动',    # 原来：'原矿产销存（吨 · 苏姆）· 全自动'
    ('资产负债表', 'H3'): '单位：美元',    # 原来：'单位：苏姆'
    ('利润表', 'D3'): '会小企02表\u3000单位：美元',    # 原来：'会小企02表\u3000单位：苏姆'
    ('现金流量表', 'A2'): '💡 直接法：【现金流水总表】每一笔按收支类别在【记账规则】里设的现金流量项目归类（收钱按「现金流量(收)」、付钱按「现金流量(付)」），账户之间转钱、换汇（项目 0）不算流入流出。月末外币调汇、转账/换汇折美元的差额单列一行「汇率变动…的影响」，所以期末现金余额正好等于资产负债表的货币资金。',    # 原来：'💡 直接法：【现金流水总表】每一笔按收支类别在【记账规则】里设的现金流量项目归类（收钱按「现金流量(收)」、付钱按「现
    ('现金流量表', 'D3'): '会小企03表\u3000单位：美元',    # 原来：'会小企03表\u3000单位：苏姆'
    ('经营报表', 'A1'): '经营报表（矿山 · 本月 / 本年累计 · 美元）',    # 原来：'经营报表（矿山 · 本月 / 本年累计 · 苏姆）'
    ('经营报表', 'B9'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B10'): '美元/吨',    # 原来：'苏姆/吨'
    ('经营报表', 'B13'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B14'): '美元/吨',    # 原来：'苏姆/吨'
    ('经营报表', 'B15'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B16'): '美元/吨',    # 原来：'苏姆/吨'
    ('经营报表', 'B17'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B19'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B20'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B22'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B23'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B24'): '美元/吨',    # 原来：'苏姆/吨'
    ('经营报表', 'B25'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B27'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B28'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B29'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B30'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B31'): '美元/月',    # 原来：'苏姆/月'
    ('经营报表', 'B34'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B35'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B36'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B37'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B38'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B39'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B41'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B43'): '美元',    # 原来：'苏姆'
    ('经营报表', 'B44'): '美元',    # 原来：'苏姆'
    ('数据校验', 'D15'): '不平说明模板公式被改动过',    # 原来：'不平说明模板公式被改动过，找我'
    ('数据校验', 'D18'): '不平一般是科目表里加了新的一级科目，报表没归进去',    # 原来：'不平一般是科目表里加了新的一级科目，报表没归进去——找我'
    ('数据校验', 'B21'): '资金账户：原币余额×月末汇率＝账上美元',    # 原来：'资金账户：原币余额×月末汇率＝账上苏姆'
    ('数据校验', 'D22'): '到【月度汇率】填报表月外币兑换美元的正确报价；USD固定为1',    # 原来：'到【月度汇率】把这个月的人民币、美元汇率填上'
    ('数据校验', 'D24'): '两边折美元的差额（包括漏记的一边）。核对清楚后再决定是否把兑换差额结转汇兑损益',    # 原来：'买卖外汇两边折苏姆的差额（＋一边没记的）。核对清楚后把【基础资料】「外币兑换差额月末转汇兑损益」改成「是」，每月自动转
    ('数据校验', 'D30'): '外币进货按进货月美元汇率记应付，付款按付款时美元汇率冲，差额看【往来余额表】P列',    # 原来：'人民币/美元进货按进货月汇率记应付，付款按付款时的汇率冲，差额留在应付账款里。到【往来余额表】P 列看是哪个供应商，做
    ('数据校验', 'D42'): '共 4000 行',    # 原来：'共 3960 行，快满了找我加'
}


# 几处格式（用户 / A075 改的）
NUMFMT = {('月度汇率', 'B6:D37'): '#,##0.0000'}               # USD / CNY / ZIG 三列汇率显示 4 位小数
FONT_SIZE = {('科目表', 'F5:F344'): 8}                         # 「外币核算」列字号 8（新加的 345、346 行没改）
# 现金流水总表第 5 行表头底色：用户 10/05 把录入列改成淡黄、自动列改成淡灰
CASH_HDR_FILL = {**{c: 'FFFFF2CC' for c in 'ABEGHIJPQR'}, **{c: 'FFF2F2F2' for c in 'CDFKLMNOSUVW'}}


def apply(wb, ctx):
    import s_base
    ref = ctx['ref']
    log = []
    ref_a075.copy_sheet(ref, wb, SH_HOME, drop_rows=(HOME_DROP_ROW,))
    build_help_report(wb, ctx)
    log.append(f'首页从原件整张照搬（第 {HOME_DROP_ROW} 行【数据导入】去掉，下面的行往上挪）；使用说明按报表本重写（录入本那份 build_help_entry）')
    log += fix_texts(wb, ctx)
    for (sh, cell), v in TEXTS.items():
        c = wb[sh][cell]
        assert c.value is None or isinstance(c.value, str), (sh, cell, c.value)     # 只改文字格，不碰公式
        c.value = v
    log.append(f'标题、提示、表头文字 {len(TEXTS)} 格改成 A075 的写法')
    log += fix_words(wb)
    log += fix_1015(wb, ctx)
    if s_base.SH_BOX not in wb.sheetnames:
        s_base.build_container(wb, ctx)
        log.append(f'新加【{s_base.SH_BOX}】（录入本，纯手填备查，{s_base.BOX_R1 - s_base.BOX_R0 + 1} 行；导入 2026 货柜单号台账 '
                   f'{len(ctx.get("containers") or [])} 行）')
    for sh, rgb in TAB_COLORS.items():
        if sh in wb.sheetnames:
            wb[sh].sheet_properties.tabColor = rgb
    for sh, d in COL_WIDTHS.items():
        for col, (w, hid) in d.items():
            wb[sh].column_dimensions[col].width = w
            wb[sh].column_dimensions[col].hidden = hid
    for sh, d in ROW_HEIGHTS.items():
        for r, h in d.items():
            wb[sh].row_dimensions[r].height = h
    for sh, d in PAGE_SETUP.items():
        ps = wb[sh].page_setup
        st = d['setup']
        ps.paperSize = int(st['paperSize']) if 'paperSize' in st else None
        ps.orientation = st.get('orientation')
        ps.scale = int(st['scale']) if 'scale' in st else None
        ps.fitToWidth = int(st['fitToWidth']) if 'fitToWidth' in st else None
        ps.fitToHeight = int(st['fitToHeight']) if 'fitToHeight' in st else None
        wb[sh].sheet_properties.pageSetUpPr.fitToPage = d['fitToPage'] == '1'
    for sh, z in ZOOM.items():
        wb[sh].sheet_view.zoomScale = int(z)
    for sh, ref_ in FILTERS.items():
        wb[sh].auto_filter.ref = ref_
    from copy import copy
    from openpyxl.styles import PatternFill
    for (sh, rg), f in NUMFMT.items():
        for row in wb[sh][rg]:
            for c in row:
                c.number_format = f
    for (sh, rg), sz in FONT_SIZE.items():
        for row in wb[sh][rg]:
            for c in row:
                ft = copy(c.font)
                ft.sz = sz
                c.font = ft
    for col, rgb in CASH_HDR_FILL.items():
        wb[SH_CASH][f'{col}{CASH_HDR}'].fill = PatternFill('solid', fgColor=rgb)
    log.append(f'标签颜色 {len(TAB_COLORS)} 张、列宽 {sum(map(len, COL_WIDTHS.values()))} 列、行高 {sum(map(len, ROW_HEIGHTS.values()))} 行、'
               f'打印设置 {len(PAGE_SETUP)} 张、显示比例 {len(ZOOM)} 张、筛选范围 {len(FILTERS)} 张照原件')
    # wps（repro 第四节）：现金流水总表 T 列 57 条 A062 鑫盛的旧批注不再照搬（A075 没有流水，那些批注对不上了）
    return log


# ───── wps：挡使用的文字和默认值（repro 第四节、optchk 第 4 节）；10/15 改成报表本的说法 ─────
HOME_DROP_ROW = 22          # 首页第 22 行：【数据导入】（A075 没有这张表，链接是断的）
HOME_EDIT = [   # (原文片段, 改成)，在首页里找唯一的一格替换
    ('★ 每一笔收付款（五个账户合一张表），自动出凭证', '★ 资金台帐：每一笔收付款（美元、人民币、FBC 三个账户一张表），自动出凭证'),
    ('所有凭证的分录总表（1.6 万行）', '所有凭证的分录总表（1.8 万行）'),
    ('选一个科目看每一笔（在表里 L3 选科目）', '选一个科目看每一笔（在表里 B3 选科目）'),
]
HOME_REPORT_MONTH = 6       # 首页报表月份默认 6 月（导进来的资金台帐到 2026-06-30；7 月起 CNY 汇率、流水还没有）


def _edit_unique(ws, old, new):
    hit = [c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and old in c.value]
    assert len(hit) == 1, (ws.title, old, [c.coordinate for c in hit])
    hit[0].value = hit[0].value.replace(old, new)
    return hit[0].coordinate


def fix_texts(wb, ctx):
    log = []
    hm = wb[SH_HOME]
    cells = [_edit_unique(hm, a, b) for a, b in HOME_EDIT]
    log.append(f'首页文字 {len(cells)} 处：五个账户→三个（资金台帐）、1.6 万行→1.8 万行、L3→B3 选科目（{",".join(cells)}）')
    b = wb[SH_BASE]
    got = [(b[f'{A_NAME}{r}'].value, b[f'{A_CODE}{r}'].value) for r in range(ACC_R0, ACC_R0 + 3)]
    assert all(c for _, c in got), got          # 10/15：对应科目由 ref_a075.apply_1015 带进来（101202 / 100105 / 100210）
    log.append('基础资料 J5:J7 对应科目：' + '、'.join(f'{n} {c}' for n, c in got) + '（照她老账，ref_a075.ACC_CODE_1015）')
    return log


# ───── 10/15：字样（plan.md 第二节第 1 条）─────
# 只改「字」：标题、提示、表头、校验提示、凭证摘要里的字；表名（产量登记、原矿产销存）不改，所以「原矿产销存」这几个字原样留着（是表名）。
# 用户录的数据（流水摘要、备注、物料单位…）一个不碰：录入表只改表头那几行，基础资料只改 A 列参数名，记账规则、科目表、月度汇率不在这里改（数据在 ref_a075）。
KEEP = ['原矿产销存']                              # 表名，不改
WORDS = [   # 按顺序替换
    ('苏姆', '美元'),
    ('卖原矿给甲方', '卖黄金给买方'), ('卖矿', '卖金'), ('甲方', '买方'),
    ('吨矿生产成本', '每克生产成本'), ('吨矿销售成本', '每克销售成本'), ('吨矿完全成本', '每克完全成本'), ('吨矿成本', '每克成本'),
    ('每吨', '每克'), ('出矿', '冶炼'), ('原矿', '黄金'), ('吨', '克'),
]
WORDS_SUMU = [('苏姆', '美元')]
_INPUT = {SH_BUY: MOD_HDR, SH_ISS: MOD_HDR, SH_PROD: MOD_HDR, SH_SALE: MOD_HDR, SH_PAY: MOD_HDR, SH_FA: MOD_HDR, SH_MAN: MOD_HDR}
_SKIP = {SH_RULE, SH_ACC, SH_RATE, SH_CASH, SH_HELP}   # 数据表（ref_a075 管）、流水（TEXTS 管表头）、使用说明（重写的）


def _swap_words(t, words):
    for i, k in enumerate(KEEP):
        t = t.replace(k, f'\x00{i}\x00')
    for a, b in words:
        t = t.replace(a, b)
    for i, k in enumerate(KEEP):
        t = t.replace(f'\x00{i}\x00', k)
    return t


def _swap_formula(f, words):
    """公式里只换双引号里的字（摘要、提示），不碰引用和函数"""
    parts = f.split('"')
    for i in range(1, len(parts), 2):
        parts[i] = _swap_words(parts[i], words)
    return '"'.join(parts)


def fix_words(wb):
    n_t = n_f = 0
    for ws in wb.worksheets:
        if ws.title in _SKIP:
            continue
        words = WORDS_SUMU if ws.title in (SH_BUY, SH_PAY, SH_FA, SH_MAN, SH_ISS) else WORDS
        top = _INPUT.get(ws.title)
        if ws.title == SH_BASE:
            top = None
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if not isinstance(v, str):
                    continue
                if v.startswith('='):
                    if ws.title == SH_BASE:
                        continue
                    nv = _swap_formula(v, words)
                    if nv != v:
                        c.value = nv
                        n_f += 1
                    continue
                if ws.title == SH_BASE and not (c.column_letter == 'A' and c.row <= 5 + len(PARAMS)):
                    continue
                if top is not None and c.row > top:
                    continue
                nv = _swap_words(v, words)
                if nv != v:
                    c.value = nv
                    n_t += 1
        for dv in ws.data_validations.dataValidation:
            for k in ('prompt', 'error', 'promptTitle'):
                t = getattr(dv, k)
                if isinstance(t, str):
                    setattr(dv, k, _swap_words(t, WORDS_SUMU))
    return [f'字样：苏姆→美元、原矿 / 吨 / 甲方 / 出矿→黄金 / 克 / 买方 / 冶炼：文字 {n_t} 格、公式里的提示和摘要 {n_f} 格（表名不改）']


# 10/15 的表头、提示（fix_words 之后写，不再过一遍替换）
TEXTS_1015 = {
    (SH_CASH, 'A1'): '资金台帐（现金流水总表）· 每一笔收钱、付钱记一行（美元记账，人民币、ZiG 按汇率折美元）',
    (SH_CASH, 'A2'): ('💡 每笔一行：填日期、选账户（美元 / 人民币 / FBC），收入、支出只填一个。类别就是你老账的一级科目（原材料、管理费用、生活支出…），'
                      '「二级/备注」写二级（公关费、PAYE、宽体车…），自动记到明细科目（对照表在【记账规则】最右边「二级明细对照」）。'
                      '人民币、ZiG 按【月度汇率】（1 美元＝多少原币）自动折美元，有实际成交汇率就填 J 列。漏记的记在最后一行下面，日期写对就行；别插行、删行。'),
    (SH_CASH, f'{K_NOTE}{CASH_HDR}'): '二级/备注',
    (SH_ISS, f'{I_NOTE}{MOD_HDR}'): '领取人/备注',
    (SH_PROD, 'A1'): '冶炼登记（产量登记 · 每次冶炼出多少克金）',
    (SH_PROD, 'A2'): ('💡 每冶炼一次记一行：日期、批号 / 来源堆号、冶炼纯金量（克）。后面掘进量、品位、贫化率是以前采矿用的，选填，不用管。'
                      '每月月底自动把当月生产成本结转成黄金库存成本（有冶炼的月份才结转），【原矿产销存】（就是黄金产销存）算出每克成本，卖出去时自动结转销售成本。'),
    (SH_PROD, f'{PD_FACE}{MOD_HDR}'): '批号/来源堆号',
    (SH_PROD, f'{PD_TON}{MOD_HDR}'): '冶炼纯金量(克)',
    (SH_PROD, f'{PD_DEV}{MOD_HDR}'): '掘进量(米)\n（选填）',
    (SH_PROD, f'{PD_GRADE}{MOD_HDR}'): '品位(%)\n（选填）',
    (SH_PROD, f'{PD_DIL}{MOD_HDR}'): '贫化率(%)\n（选填）',
    (SH_SALE, 'A1'): '黄金销售（销售结算 · 按克卖给买方）',
    (SH_SALE, 'A2'): ('💡 卖一次黄金记一行：日期、买方（私人 / 银行 / 银行代理，也可以直接写名字）、出售纯金量（克）、单价（美元/克），金额自动算；品名空着就是黄金。'
                      '不交销项税（【基础资料】增值税率 0、单价不含税）。自动凭证：借 应收账款-黄金销售，贷 主营业务收入-黄金销售；月底按每克成本自动结转销售成本。'
                      '收到钱在资金台帐类别选「主营业务收入」，自动冲应收（业务模块启用日以前的直接记收入）。'),
    (SH_SALE, f'{S_CUS}{MOD_HDR}'): '买方',
    (SH_SALE, f'{S_ITEM}{MOD_HDR}'): '品名\n(空着＝黄金)',
    (SH_SALE, f'{S_TON}{MOD_HDR}'): '出售纯金量\n(克)',
    (SH_SALE, f'{S_PRICE}{MOD_HDR}'): '单价\n(美元/克)',
    (SH_SALE, f'{S_UC}{MOD_HDR}'): '每克成本\n(自动)',
    (SH_SALE, f'{S_COGS}{MOD_HDR}'): '销售成本\n(自动)',
    (SH_ORE, 'A1'): '黄金产销存（克 · 美元）· 全自动',
    (SH_OPS, 'E39'): '采矿权、土地使用权',
    (SH_ARAP, 'M4'): '外币应付\n津币 ZiG(原币)',
    (SH_ACC, 'A2'): ('💡 编码按 4-2-2-2 分级。末级科目“期初余额”在建账起始日登记（借正贷负）。现在只填了三个资金账户（美元 100105、人民币 101202、FBC 100210），'
                     '2025 年底其他科目的余额还没有，差额先放在 310499「期初差额待核」（科目余额表、资产负债表才平）；拿到 2025 年底的余额表后补上，再把 310499 清成 0。'
                     '要加明细科目往下面空行填，灰色列自动计算，不要删除预留公式。'),
    (SH_HOME, 'A2'): ('💡 这本是报表本：只看不录。数都从同一个文件夹里的「A075-1_录入本」取（先把录入本存盘，再打开这本点「启用内容 / 更新」）。'
                      '黄格子改报表年度、月份，所有账簿、报表、凭证打印都跟着变。1–6 月的流水是从你的资金台帐导进来的；期初只补了三个资金账户'
                      '（美元 100105、人民币 101202、FBC 100210），2025 年底其他科目的余额还没填，差额先放在 310499「期初差额待核」。'),
    (SH_HOME, 'B10'): '本月冶炼（克）',
    (SH_HOME, 'B16'): '① 设置层 —— 在「A075-1_录入本」里改（这本里是隐藏的取数表）',
    (SH_HOME, 'B21'): '② 录入层 —— 在「A075-1_录入本」里录（这本里是隐藏的取数表）',
    (SH_HOME, 'F22'): '★ 资金台帐：每一笔收付款（美元、人民币、FBC 三个账户一张表），自动出凭证',
    (SH_HOME, 'F23'): '购进：炸药、氰化钠、盐酸、柴油、配件、建材、食材…',
    (SH_HOME, 'F25'): '冶炼登记：每次冶炼出多少克金（批号 / 来源堆号）',
    (SH_HOME, 'F26'): '黄金销售：买方、克数、单价（美元/克），自动出凭证、结转成本',
    (SH_HOME, 'B50'): '① 平时在「A075-1_录入本」里录：资金台帐（现金流水总表）、购进（采购入库）、黄金销售（销售结算）、领用、冶炼登记、工资、设备卡片。',
    (SH_HOME, 'B51'): '② 看账前：先把录入本存盘，再打开这本点「启用内容」或「更新」；看上面那行「数据来自录入本：流水 N 笔，最后一笔 …」对不对。',
    (SH_HOME, 'B52'): '③ 月底：在录入本【月度汇率】填当月汇率（1 美元＝多少人民币 / ZiG）；工资、新买的设备也在录入本登记。',
    (SH_HOME, 'B54'): '⑤ 这本【首页】把报表月份改成这个月 → 看【数据校验】有没有 ✗ → 看【凭证汇总】【科目余额表】→ 打印【记账凭证】和三张报表。',
    (SH_HOME, 'B55'): '⑥ ✗ 要回录入本去改（这本里改了不算，下次更新会被盖掉）；⚠ 只是提醒。',
}
HOME_INPUT_ROWS = (17, 18, 19, 20, 22, 23, 24, 25, 26, 27, 28, 29)     # 首页导航里设置层、录入层那几行（E 列「你要填」→「在录入本填」）


def fix_1015(wb, ctx):
    log = []
    # 表头、提示
    for (sh, cell), v in TEXTS_1015.items():
        c = wb[sh][cell]
        assert c.value is None or (isinstance(c.value, str) and not c.value.startswith('=')), (sh, cell, c.value)
        c.value = v
    hm = wb[SH_HOME]
    for r in HOME_INPUT_ROWS:
        assert hm[f'E{r}'].value in ('你要填', '每月填'), (r, hm[f'E{r}'].value)
        hm[f'E{r}'] = '在录入本填'
    hm.row_dimensions[2].height = 64
    hm['F4'] = HOME_REPORT_MONTH
    log.append(f'表头 / 提示 {len(TEXTS_1015)} 格（黄金销售、冶炼登记、二级/备注、领取人/备注、首页报表本说法）；首页报表月份默认 {HOME_REPORT_MONTH} 月')
    # 往来余额表：外币应付第二列从「美元」改成津币 ZiG（本位币就是美元，美元没有「外币应付」）
    ws = wb[SH_ARAP]
    n = 0
    for row in ws.iter_rows(min_row=5, min_col=CI('M'), max_col=CI('O')):
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and '"USD"' in c.value:
                c.value = c.value.replace('"USD"', '"ZIG"')
                n += 1
    assert n > 0
    log.append(f'往来余额表 M～O 外币应付改按津币 ZiG 算（公式里的 "USD"→"ZIG"，{n} 格）')
    for sh, cell, a, b in [(SH_ARAP, 'A2', '人民币、美元进货', '人民币、津币 ZiG 进货'),
                           (SH_BUY, 'A2', '外币（人民币/美元）选币种', '外币（人民币 / 津币 ZiG）选币种'),
                           (SH_MEND, 'A2', '② 调汇：人民币、美元账户', '② 调汇：人民币、FBC（ZiG）账户')]:
        v = wb[sh][cell].value
        assert isinstance(v, str) and v.count(a) == 1, (sh, cell, a)
        wb[sh][cell] = v.replace(a, b)
    # 销售结算：品名下拉「黄金」（空着凭证摘要也写黄金）
    add_list_dv(wb[SH_SALE], f'{S_ITEM}{MOD_R0}:{S_ITEM}{SALE_R1}', '"黄金"', '空着就是黄金', stop=False)
    return log


# ───── 10/15：首页导航加 s_extra 的 5 张新报表（build.make_onebook 在 s_extra.build 之后调）─────
HOME_NAV_EXTRA = [   # (表名, 你要做什么, 做什么用)
    ('发生情况报表', '自动·选月份', '你月报的样子：资金来源、资金占用、资金结余（美金 / 人民币 / 津币），期初、本期、期末'),
    ('记账底稿', '自动·选账户月份', '你底稿帐的样子：选账户、月份，逐笔列凭证号、借贷科目、折美金、原币'),
    ('物资出入库台账', '自动·选物资', '选一种物资看每一笔进、出、结余（进＝采购入库，出＝领用出库）'),
    ('黄金总记录', '自动·选年份', '每月冶炼、出售克数和差额；逐笔卖金：克数、单价、累计收入、折人民币'),
    ('收支分类报表', '自动·选年份', '你的收支类别 × 1～12 月（美元），收入、支出分开'),
]
HOME_NAV_HDR = '⑤ 手工表样子的报表 —— 全自动（黄格子选月份 / 账户 / 物资 / 年份，空着就跟上面的报表月份）'


def home_nav_extra(wb, items=None, at=None):
    """首页导航④后面加一节「⑤ 手工表样子的报表」（一行表头＋每张表一行，样子照第 36 行 / 第 45 行），下面「每月怎么做」整块往下挪。
    at＝新一节表头放第几行（默认：导航最后一个链接下面那一行）。只加 wb 里有的表。返回日志行"""
    from copy import copy
    from openpyxl.worksheet.hyperlink import Hyperlink
    items = [x for x in (items or HOME_NAV_EXTRA) if x[0] in wb.sheetnames]
    if not items:
        return []
    ws = wb[SH_HOME]
    links = sorted(c.row for row in ws.iter_rows() for c in row if c.hyperlink is not None and c.column == 3)
    at = at or links[-1] + 1
    sec = next(r for r in range(at - 1, 0, -1) if isinstance(ws[f'B{r}'].value, str) and ws[f'B{r}'].value.startswith('④'))
    tpl = links[-1] - 2 if links[-1] - 2 > sec else links[-1]        # 普通一行（不是「先看这个」那种加粗行）的样子
    n = len(items) + 1
    last = ws.max_row
    # 下面的块往下挪 n 行：合并格、行高跟着挪（openpyxl move_range 只挪格子）
    mv = [mr for mr in list(ws.merged_cells.ranges) if mr.min_row >= at]
    for mr in mv:
        ws.unmerge_cells(str(mr))
    hts = {r: ws.row_dimensions[r].height for r in range(at, last + 1)}
    assert not any(c.hyperlink is not None for row in ws.iter_rows(min_row=at) for c in row), '要挪的行里有链接'
    ws.move_range(f'A{at}:{CL(ws.max_column)}{last}', rows=n)
    for mr in mv:
        ws.merge_cells(start_row=mr.min_row + n, start_column=mr.min_col, end_row=mr.max_row + n, end_column=mr.max_col)
    for r in range(at, last + n + 1):
        ws.row_dimensions[r].height = hts.get(r - n)

    def clone(src_r, dst_r):
        for ci in range(1, ws.max_column + 1):
            a, b = ws.cell(src_r, ci), ws.cell(dst_r, ci)
            if a.has_style:
                b.font, b.fill, b.border = copy(a.font), copy(a.fill), copy(a.border)
                b.alignment, b.number_format = copy(a.alignment), a.number_format
        ws.row_dimensions[dst_r].height = ws.row_dimensions[src_r].height
        for mr in list(ws.merged_cells.ranges):
            if mr.min_row == mr.max_row == src_r:
                ws.merge_cells(start_row=dst_r, start_column=mr.min_col, end_row=dst_r, end_column=mr.max_col)

    clone(sec, at)
    ws[f'B{at}'] = HOME_NAV_HDR
    for i, (sh, todo, desc) in enumerate(items):
        r = at + 1 + i
        clone(tpl, r)
        ws[f'C{r}'] = sh
        ws[f'E{r}'] = todo
        ws[f'F{r}'] = desc
        h = copy(ws[f'C{tpl}'].hyperlink)
        h.location, h.ref, h.id, h.target = f"'{sh}'!A1", f'C{r}', None, None
        h.display = sh if h.display else None
        h.tooltip = None
        ws[f'C{r}'].hyperlink = h
    return [f'首页导航加「⑤ 手工表样子的报表」{len(items)} 张（第 {at}～{at + len(items)} 行），下面「每月怎么做」往下挪 {n} 行']


# ───── 10/15：使用说明两份（大白话、短）─────
HELP_REPORT = [   # (None, 小标题) 或 (编号自动, 文字)
    (None, '这本是干什么的'),
    '这本是「报表本」：只看不录。凭证、账簿、三张报表、经营报表……都从同一个文件夹里的「A075-1_录入本」自动取数。',
    '这本里只有「看哪个月、哪个科目、哪个账户」这种黄格子（【首页】的报表年度、月份，【明细账】的科目……）。别的格子别填：填了也回不到录入本，下次更新就被盖掉。',
    (None, '怎么看到最新的数'),
    '① 先把录入本存盘（存完可以关掉）。② 打开这本，上面出现黄条「安全警告 / 已禁用外部链接」就点「启用内容」（WPS 点「更新」或「启用」）。'
    '③ 看【首页】那一行「数据来自录入本：流水 N 笔，最后一笔 …」，笔数、日期对得上就是最新的。',
    '两本放在同一个文件夹里，文件名别改（A075-1_录入本.xlsx、A075-2_报表本.xlsx）。',
    '改了录入本的名字、挪了地方、从微信另存多出「(1)」：这本还是旧数。点「数据 → 编辑链接 → 更改源」，选现在那本录入本，再点「更新」。',
    '两本别同时开着录（同时开着，录入本每录一格都要等这本算一遍，又会卡）。',
    (None, '每个月'),
    '【首页】黄格子改报表年度、月份 → 先看【数据校验】：✗ 要回录入本去改，⚠ 看一眼就行 → 再看【凭证汇总】【科目余额表】【资产负债表】【利润表】。',
    '打印凭证：【凭证汇总】看「A4 张数」→【记账凭证】打印，页码 1 到 A4 张数（A4 纵向，一张纸上下两页凭证）。月末结转、年末结转的凭证排在当月最后几号。',
    '月末结转不用做：折旧、调汇、生产成本转黄金、结转销售成本、结转损益全自动，日期是月末最后一天；12 月 31 日自动把本年利润转进未分配利润。',
    (None, '要知道的'),
    '1–6 月的流水是从你的资金台帐导进来的。期初只补了三个资金账户（美元 100105、人民币 101202、FBC 100210），2025 年底其他科目的余额还没填，'
    '差额先放在【科目表】310499「期初差额待核」，所以科目余额表、资产负债表是平的；拿到 2025 年底的余额表后在录入本【科目表】M 列补上，把 310499 清成 0。',
    '利润是「按付款记」的：9 月 1 日（业务模块启用日）以前买材料付的钱直接算成本，没有存货、没有折旧；要准，以后走采购入库、领用、设备卡片。',
    '你手工表样子的 5 张：【发生情况报表】（月报）、【记账底稿】、【物资出入库台账】、【黄金总记录】、【收支分类报表】。'
    '表头黄格子选月份 / 账户 / 物资 / 年份，空着就跟【首页】的报表月份；清单按最多的行数留了位置，打印只印有数的那几页。',
    '数字不对先看【数据校验】；WPS 里建议开「自动计算」，打开这本要等它算完。',
]
HELP_ENTRY = [
    (None, '这本是干什么的'),
    '这本是「录入本」：只用来录账。资金台帐（现金流水总表）、购进（采购入库）、黄金销售（销售结算）、领用、冶炼登记、工资、设备卡片、手工凭证、货柜登记都在这本里录。'
    '凭证、账簿、报表在另一本「A075-2_报表本」里自动出。',
    '录的时候只开这一本，别同时开着报表本（同时开着会卡）。',
    (None, '怎么录'),
    '淡黄、白色格子手填；灰格子自动算，别往里写。清单里没有的（账户、类别、物资、买方 / 供应商、科目），先去【基础资料】【记账规则】【科目表】往下面空行加一行，再回来录。',
    '资金台帐：一笔一行。填日期、选账户（美元 / 人民币 / FBC）、选类别（就是你老账的一级科目：原材料、管理费用、生活支出……），收入、支出只填一个。'
    '「二级/备注」写二级（公关费、PAYE、宽体车……），会自动记到明细科目（对照表在【记账规则】最右边「二级明细对照」，没有的可以往下加）。',
    '人民币、FBC（ZiG）的钱自动按【月度汇率】折美元：每个月在【月度汇率】填好当月汇率（1 美元＝多少人民币 / ZiG）。',
    '账户之间调钱、换汇：转出、转入两笔都要记，类别都选「购汇/结汇」（同一种钱转账选「内部转账」）。',
    '卖黄金：【销售结算】一次一行，填日期、买方、克数、单价（美元/克）；收到钱在资金台帐类别选「主营业务收入」。'
    '冶炼：【产量登记】每次一行，填日期、批号 / 来源堆号、冶炼纯金量（克）。',
    '进货：【采购入库】一次一个物资一行；付钱在资金台帐选「原材料」「柴油」「氰化钠」……。领料：【领用出库】选物资、用途（堆淋、氰化厂……），领取人写在最后一列。',
    (None, '千万别做'),
    '别插行、删行。漏记的一笔，记在最后一行下面就行，日期写对，报表按日期算，不用排在前面。要删一笔：选中那一行的格子按 Delete 清空。',
    '别在表头上面插行，别插列、删列，别改表名，别改这本的文件名（报表本靠名字找它）。',
    (None, '录完'),
    '看一眼【录入检查】（【首页】也显示）：✗ 一定要改，⚠ 看一眼就行。',
    '存盘（Ctrl+S）。存好了再去开报表本看账。',
]


def _help(wb, ctx, items, ttl):
    idx = wb.sheetnames.index(SH_HELP) if SH_HELP in wb.sheetnames else None
    if idx is not None:
        del wb[SH_HELP]
    ws = wb.create_sheet(SH_HELP, idx)
    widths(ws, {'A': 6, 'B': 112})
    title(ws, ttl, 'B', C_HOME, h2=8)
    ws['A2'] = None
    r, k = 3, 0
    for x in items:
        if isinstance(x, tuple):
            ws.merge_cells(f'A{r}:B{r}')
            put(ws, f'A{r}', x[1], F_SEC, FILL_SUBH, align=AL)
            put(ws, f'B{r}', None, fill_=FILL_SUBH)
            ws.row_dimensions[r].height = 24
        else:
            k += 1
            put(ws, f'A{r}', k, F_TXTB, align=AC)
            put(ws, f'B{r}', x, F_TXT, align=ALW)
            ws.row_dimensions[r].height = 18 * max(1, -(-len(x) // 58)) + 6
        r += 1
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.paperSize = 9
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.tabColor = C_HOME[2:]
    return ws


def build_help_report(wb, ctx=None):
    """【使用说明】报表本版（一本版里也是这份）：先存录入本、打开报表本点更新 / 启用内容、两本放同一文件夹别改名、改了名用 数据→编辑链接→更改源"""
    return _help(wb, ctx, HELP_REPORT, '使用说明（报表本：只看不录）')


def build_help_entry(wb, ctx=None):
    """【使用说明】录入本版（split.make_entry 里调，换掉一本版那份）：怎么录、别插行（漏记写在最后）、录的时候只开录入本、存盘"""
    return _help(wb, ctx, HELP_ENTRY, '使用说明（录入本：平时录账只开这一本）')
