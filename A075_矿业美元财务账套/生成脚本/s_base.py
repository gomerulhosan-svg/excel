# -*- coding: utf-8 -*-
"""【基础资料】【记账规则】【月度汇率】【科目表】+ 整本册子用的定义名称。"""
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

from common import *

# 10/15：【记账规则】最右边「二级明细对照」（类别＋二级 → 科目）。现金流水 Y 列（s_cash.raw_formula）先查它：
# 键＝类别&"|"&二级（流水 E&"|"&R），对上了就记这里的科目，没对上照收支类别记。名称 二级_键 / 二级_科目（define_names）
SUB_R0, SUB_R1 = 5, 84             # 80 行
SB_CAT, SB_SUB, SB_ACC, SB_NOTE, SB_KEY = 'AM AN AO AP AQ'.split()     # 类别 / 二级 / 记到科目 / 说明 / 键（隐藏，公式）
# 10/15：【货柜登记】（录入本，纯手填备查，不记账）
SH_BOX = '货柜登记'
BOX_HDR, BOX_R0, BOX_R1 = 3, 4, 303                                 # 300 行
BOX_COLS = [('A', '签收日期', 11), ('B', '柜号', 16), ('C', '提单号', 20), ('D', '封条号', 14), ('E', '货物/供应商', 28), ('F', '备注', 36)]


def code_of(x):
    """单元格里可能是「160101」也可能是下拉选的「160101 矿用车辆…」：取空格前的编码"""
    return f'IF(ISNUMBER(FIND(" ",{x}&"")),LEFT({x},FIND(" ",{x})-1),{x}&"")'


def _block_title(ws, c0, c1, text, color, row=3):
    ws.merge_cells(f'{c0}{row}:{c1}{row}')
    put(ws, f'{c0}{row}', text, F_HDR, fill(color), align=AC)
    for cc in range(CI(c0) + 1, CI(c1) + 1):
        put(ws, f'{CL(cc)}{row}', None, fill_=fill(color))


# ───────────────────────── 基础资料 ─────────────────────────
def build_base(wb, ctx):
    ws = wb.create_sheet(SH_BASE)
    widths(ws, {'A': 26, 'B': 30, 'C': 26, A_SEQ: 5, A_NAME: 24, A_BANK: 10, A_CUR: 7, A_OPEN: 15, A_OPENB: 16,
                A_CODE: 10, A_NOWO: 16, A_NOWB: 18, A_NOTE: 20, 'N': 2, PJ_SEQ: 5, PJ_NAME: 14, 'Q': 2,
                CP_SEQ: 5, CP_NAME: 20, CP_TYPE: 8, CP_CAP: 11, CP_NOTE: 22, 'W': 2,
                M_SEQ: 5, M_NAME: 16, M_SPEC: 10, M_UNIT: 6, M_CAT: 11, M_ACC: 10, M_STOCK: 7, M_Q0: 10,
                M_P0: 11, M_A0: 13, M_NOTE: 16})
    title(ws, '基础资料（参数 · 资金账户 · 项目 · 往来单位 · 物料档案）', M_NOTE, C_BASE,
          '💡 淡黄格子手填，灰格子自动。这里的清单就是各录入表的下拉内容——清单里没有的先在这里加一行，再去录。'
          '不要插入或删除行，直接往下面的空行里填。收支类别、物料类别、领用用途这些「记到哪个科目」的规则在【记账规则】。')
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 34

    # 参数 A:B
    _block_title(ws, 'A', 'B', '参数', C_BASE)
    header(ws, 4, [('A', '项目'), ('B', '值')], 'FF808080')
    fixes = ctx['fx']
    vals = ctx['params']            # A075：参数值全部照原件（A062 这里是写死的鑫盛参数）
    fmts = {'start': MONTH, 'golive': DATE, 'vat': '0%', 'fxstart': MONTH, **{k: '@' for k in ACC_PARAMS}}
    for key, name, _ in PARAMS:
        r = P_ROW[key]
        put(ws, f'A{r}', name, F_TXTB, FILL_SUBH, align=AL)
        put(ws, f'B{r}', vals[key], F_IN, FILL_IN, fmts.get(key), AC)
    add_list_dv(ws, f'B{P_ROW["vatin"]}', '"是,否"')
    add_list_dv(ws, f'B{P_ROW["vnoprint"]}', '"否,是"', '「否」＝凭证上的号留空，手写；自动编的号在打印区右边的辅助列')
    add_list_dv(ws, f'B{P_ROW["fxclose"]}', '"否,是"', '「是」＝从下一格那个月起，每月把「外币兑换待对冲」的余额转进财务费用-汇兑损益')
    add_list_dv(ws, f'B{P_ROW[ACC_PARAMS[0]]}:B{P_ROW[ACC_PARAMS[-1]]}', '=科目选择', '填编码或从下拉选', stop=False)
    for k in ACC_PARAMS:
        put(ws, f'C{P_ROW[k]}', f'=IFERROR(INDEX({acc(AC_FULL)},MATCH({code_of(f"B{P_ROW[k]}")},{acc(AC_CODE)},0)),"✗ 科目表里没有")',
            F_NOTE, align=AL, border=False)
    r = P_ROW[ACC_PARAMS[-1]] + 2
    notes = [
        '· 建账起始月份：流水从这个月开始；月度汇率、月末结转都从这个月往后排 32 个月。',
        '· 业务模块启用日：这天以前的流水没有采购入库、工资计提、销售结算，按你 8 月报表的现金口径直接记成本费用；'
        '这天起「采购付款」冲应付、「应付职工薪酬」冲计提、「销售回款」冲应收。',
        '· 增值税率：销售结算自动算销项税；采购的进项税在采购入库里填。',
        '· 凭证号是否打印：默认「否」，打出来的凭证号留空手写，自动编的号在打印区右边。',
        '· 外币兑换差额：选「是」以后，从「兑换差额从哪个月开始转」那个月起，每月月末把 122105 的余额转汇兑损益；'
        '开始那个月一次把以前累计的都转进去，以前的月份不动。',
        '· 下面几个「…科目」是采购入库、销售结算、月末结转自动出凭证用的科目，右边灰字是科目全称，改了马上生效。',
    ]
    for i, t in enumerate(notes):
        ws.merge_cells(f'A{r + i}:B{r + i}')
        put(ws, f'A{r + i}', t, F_NOTE, align=ALW, border=False)
        ws.row_dimensions[r + i].height = 30

    # 资金账户
    _block_title(ws, A_SEQ, A_NOTE, '资金账户（现金流水「账户名称」下拉）', C_CASH)
    header(ws, 4, [(A_SEQ, '序号'), (A_NAME, '账户名称'), (A_BANK, '银行'), (A_CUR, '币种'), (A_OPEN, '期初余额(原币)'),
                   (A_OPENB, f'期初余额(折{BASE_NAME})'), (A_CODE, '对应科目'), (A_NOWO, '当前余额(原币)'),
                   (A_NOWB, f'当前余额(账面{BASE_NAME})'), (A_NOTE, '备注（资金账户与会计科目一致）')], C_CASH)
    accs = ctx['accounts']
    for i in range(ACC_R1 - ACC_R0 + 1):
        r = ACC_R0 + i
        a = accs[i] if i < len(accs) else {}
        put(ws, f'{A_SEQ}{r}', f'=IF({A_NAME}{r}="","",ROW()-{ACC_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{A_NAME}{r}', a.get('name'), F_IN, FILL_IN, align=AC)
        put(ws, f'{A_BANK}{r}', a.get('bank'), F_IN, FILL_IN, align=AC)
        put(ws, f'{A_CUR}{r}', a.get('cur'), F_IN, FILL_IN, align=AC)
        put(ws, f'{A_OPEN}{r}', a.get('open') if a else None, F_IN, FILL_IN, MONEY, AR_)
        put(ws, f'{A_OPENB}{r}', a.get('openb') if a else None, F_IN, FILL_IN, MONEY, AR_)
        c = put(ws, f'{A_CODE}{r}', a.get('code'), F_IN, FILL_IN, '@', AC)
        put(ws, f'{A_NOWO}{r}', f'=IF({A_NAME}{r}="","",N({A_OPEN}{r})+SUMIF({cash_u(K_ACC)},{A_NAME}{r},{cash_u(K_NET)}))',
            F_AUTO, FILL_AUTO, MONEY, AR_)
        put(ws, f'{A_NOWB}{r}', f'=IF({A_NAME}{r}="","",N({A_OPENB}{r})+SUMIF({cash_u(K_ACC)},{A_NAME}{r},{cash_u(K_UZS)}))',
            F_AUTO, FILL_AUTO, MONEY, AR_)
        put(ws, f'{A_NOTE}{r}', a.get('note'), F_IN, FILL_IN, align=AL)
    add_list_dv(ws, f'{A_CUR}{ACC_R0}:{A_CUR}{ACC_R1}', f'={q(SH_RATE)}!$B${RATE_HDR}:$I${RATE_HDR}')
    add_list_dv(ws, f'{A_CODE}{ACC_R0}:{A_CODE}{ACC_R1}', '=科目选择', '选 1001/1002 下面的明细科目', stop=False)

    # 项目/业务
    _block_title(ws, PJ_SEQ, PJ_NAME, '项目/业务', 'FF548235')
    header(ws, 4, [(PJ_SEQ, '序号'), (PJ_NAME, '项目/业务')], 'FF548235')
    prj = ctx['projects']
    for i in range(PRJ_R1 - PRJ_R0 + 1):
        r = PRJ_R0 + i
        put(ws, f'{PJ_SEQ}{r}', f'=IF({PJ_NAME}{r}="","",ROW()-{PRJ_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{PJ_NAME}{r}', prj[i] if i < len(prj) else None, F_IN, FILL_IN, align=AC)

    # 往来单位
    _block_title(ws, CP_SEQ, CP_NOTE, '往来单位（客户/供应商/股东/员工…）', 'FF7030A0')
    header(ws, 4, [(CP_SEQ, '序号'), (CP_NAME, '往来单位'), (CP_TYPE, '类型'), (CP_CAP, '实收资本明细\n(股东才填)'),
                   (CP_NOTE, '备注')], 'FF7030A0')
    cps = ctx['counterparties']
    for i in range(CP_R1 - CP_R0 + 1):
        r = CP_R0 + i
        c = cps[i] if i < len(cps) else {}
        put(ws, f'{CP_SEQ}{r}', f'=IF({CP_NAME}{r}="","",ROW()-{CP_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{CP_NAME}{r}', c.get('name'), F_IN, FILL_IN, align=AC)
        put(ws, f'{CP_TYPE}{r}', c.get('type'), F_IN, FILL_IN, align=AC)
        put(ws, f'{CP_CAP}{r}', c.get('cap'), F_IN, FILL_IN, '@', AC)
        put(ws, f'{CP_NOTE}{r}', c.get('note'), F_IN, FILL_IN, align=AL)
    add_list_dv(ws, f'{CP_TYPE}{CP_R0}:{CP_TYPE}{CP_R1}', '"客户,供应商,股东,员工,其他"')

    # 物料档案
    _block_title(ws, M_SEQ, M_NOTE, '物料档案（采购入库、领用出库的「物料」下拉）', 'FF806000')
    header(ws, 4, [(M_SEQ, '序号'), (M_NAME, '物料名称'), (M_SPEC, '规格'), (M_UNIT, '单位'), (M_CAT, '物料类别'),
                   (M_ACC, '入账科目\n(自动)'), (M_STOCK, '库存\n管理'), (M_Q0, '期初数量'), (M_P0, '期初单价'),
                   (M_A0, '期初金额'), (M_NOTE, '备注')], 'FF806000')
    mats = fixes.MATERIALS          # (名称, 规格, 单位, 类别, 期初数量, 期初单价, 备注)，照原件
    mc = lambda col: rng(SH_RULE, col, MC_R0, MC_R1)
    for i in range(MAT_R1 - MAT_R0 + 1):
        r = MAT_R0 + i
        m = mats[i] if i < len(mats) else (None,) * 7
        put(ws, f'{M_SEQ}{r}', f'=IF({M_NAME}{r}="","",ROW()-{MAT_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{M_NAME}{r}', m[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{M_SPEC}{r}', m[1] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{M_UNIT}{r}', m[2], F_IN, FILL_IN, align=AC)
        put(ws, f'{M_CAT}{r}', m[3], F_IN, FILL_IN, align=AC)
        put(ws, f'{M_ACC}{r}', f'=IF({M_CAT}{r}="","",IFERROR(INDEX({mc(MC_ACC)},MATCH({M_CAT}{r},{mc(MC_NAME)},0))&"",""))',
            F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{M_STOCK}{r}', f'=IF({M_CAT}{r}="","",IFERROR(INDEX({mc(MC_STOCK)},MATCH({M_CAT}{r},{mc(MC_NAME)},0))&"",""))',
            F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{M_Q0}{r}', m[4], F_IN, FILL_IN, QTY, AR_)
        put(ws, f'{M_P0}{r}', m[5], F_IN, FILL_IN, PRICE, AR_)
        put(ws, f'{M_A0}{r}', f'=IF(N({M_Q0}{r})=0,"",ROUND({M_Q0}{r}*N({M_P0}{r}),2))', F_AUTO, FILL_AUTO, MONEY, AR_)
        put(ws, f'{M_NOTE}{r}', m[6], F_IN, FILL_IN, align=AL)
    add_list_dv(ws, f'{M_CAT}{MAT_R0}:{M_CAT}{MAT_R1}', '=物料类别列表')
    ws.freeze_panes = 'A5'
    return ws


# ───────────────────────── 记账规则 ─────────────────────────
def build_rules(wb, ctx):
    fixes = ctx['fx']
    ws = wb.create_sheet(SH_RULE)
    widths(ws, {RC_SEQ: 5, RC_NAME: 16, RC_ATTR: 8, RC_IN: 12, RC_OUT: 12, RC_HIST: 13, RC_CFIN: 8, RC_CFOUT: 8,
                RC_NOTE: 46, 'J': 2, MC_NAME: 12, MC_ACC: 10, MC_STOCK: 7, MC_PROD: 12, MC_NOTE: 22, 'P': 2,
                US_NAME: 12, US_ACC: 12, US_NOTE: 30, 'T': 2, DP_NAME: 10, DP_WAGE: 11, DP_DEP: 11, DP_NOTE: 22,
                'Y': 2, PT_NAME: 8, PT_ACC: 10, 'AB': 2, FC_NAME: 12, FC_ACC: 10, FC_YEARS: 7, FC_SALV: 7,
                FC_NOTE: 14, 'AH': 2, CFI_NO: 5, CFI_NAME: 44, CFI_SIDE: 5, 'AL': 2, SB_CAT: 14, SB_SUB: 16,
                SB_ACC: 12, SB_NOTE: 34, SB_KEY: 10})
    title(ws, '记账规则（每类业务记到哪个科目 —— 改这里，凭证和报表跟着变）', CFI_SIDE, C_BASE,
          '💡 科目格子填科目编码（也可以从下拉选「编码 名称」）。现金流水每一笔：选了哪个收支类别，就按这里「收到钱时 / 付出钱时」'
          '的科目记对方科目（本账户那一边按【基础资料·资金账户】的对应科目）；业务模块启用日以前的，「启用日前」一栏有填就用它（按你 8 月报表的口径）。'
          '「@资本」＝按往来单位找【基础资料·往来单位】里的实收资本明细。现金流量项目填会小企03 表的行次，0＝账户之间转钱，不算现金流。'
          '最右边「二级明细对照」：流水「二级/备注」写了二级（公关费、PAYE、宽体车…）就按它记到明细科目，对不上的照收支类别记。')
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 44

    _block_title(ws, RC_SEQ, RC_NOTE, '收支类别 → 对方科目（现金流水用）', C_CASH)
    header(ws, 4, [(RC_SEQ, '序号'), (RC_NAME, '收支类别'), (RC_ATTR, '属性'), (RC_IN, '收到钱时\n对方科目'),
                   (RC_OUT, '付出钱时\n对方科目'), (RC_HIST, '启用日前\n对方科目'), (RC_CFIN, '现金流量\n(收)'),
                   (RC_CFOUT, '现金流量\n(付)'), (RC_NOTE, '说明')], C_CASH, height=44)
    rules = ctx['rules']
    for i in range(RC_R1 - RC_R0 + 1):
        r = RC_R0 + i
        x = rules[i] if i < len(rules) else (None,) * 8
        put(ws, f'{RC_SEQ}{r}', f'=IF({RC_NAME}{r}="","",ROW()-{RC_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{RC_NAME}{r}', x[0], F_IN, FILL_IN, align=AL)
        put(ws, f'{RC_ATTR}{r}', x[1], F_IN, FILL_IN, align=AC)
        for col, v in ((RC_IN, x[2]), (RC_OUT, x[3]), (RC_HIST, x[4])):
            put(ws, f'{col}{r}', v, F_IN, FILL_IN, '@', AC)
        put(ws, f'{RC_CFIN}{r}', x[5], F_IN, FILL_IN, '0', AC)
        put(ws, f'{RC_CFOUT}{r}', x[6], F_IN, FILL_IN, '0', AC)
        put(ws, f'{RC_NOTE}{r}', x[7] or None, F_NOTE, FILL_IN, align=ALW)
        ws.row_dimensions[r].height = 26 if (x[7] and len(x[7]) > 26) else 16
    add_list_dv(ws, f'{RC_ATTR}{RC_R0}:{RC_ATTR}{RC_R1}', '"收入,支出,内部转账"')
    add_list_dv(ws, f'{RC_IN}{RC_R0}:{RC_HIST}{RC_R1}', '=科目选择', '填编码或从下拉选', stop=False)

    _block_title(ws, MC_NAME, MC_NOTE, '物料类别 → 存货/资产科目', 'FF806000')
    header(ws, 4, [(MC_NAME, '物料类别'), (MC_ACC, '入账科目'), (MC_STOCK, '库存\n管理'), (MC_PROD, '采矿生产\n领用时计入'),
                   (MC_NOTE, '说明')], 'FF806000', height=44)
    for i in range(MC_R1 - MC_R0 + 1):
        r = MC_R0 + i
        x = fixes.MAT_CATS[i] if i < len(fixes.MAT_CATS) else (None,) * 5
        put(ws, f'{MC_NAME}{r}', x[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{MC_ACC}{r}', x[1], F_IN, FILL_IN, '@', AC)
        put(ws, f'{MC_STOCK}{r}', x[2], F_IN, FILL_IN, align=AC)
        put(ws, f'{MC_PROD}{r}', x[3], F_IN, FILL_IN, '@', AC)
        put(ws, f'{MC_NOTE}{r}', x[4] or None, F_NOTE, FILL_IN, align=ALW)
    add_list_dv(ws, f'{MC_STOCK}{MC_R0}:{MC_STOCK}{MC_R1}', '"是,否"', '「是」＝进库、领用时出库；「否」＝买来直接记入账科目（设备、食材）')

    _block_title(ws, US_NAME, US_NOTE, '领用用途 → 计入科目', 'FF8064A2')
    header(ws, 4, [(US_NAME, '用途'), (US_ACC, '计入科目'), (US_NOTE, '说明')], 'FF8064A2', height=44)
    for i in range(US_R1 - US_R0 + 1):
        r = US_R0 + i
        x = fixes.USES[i] if i < len(fixes.USES) else (None,) * 3
        put(ws, f'{US_NAME}{r}', x[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{US_ACC}{r}', x[1], F_IN, FILL_IN, '@', AC)
        put(ws, f'{US_NOTE}{r}', x[2] or None, F_NOTE, FILL_IN, align=ALW)

    _block_title(ws, DP_NAME, DP_NOTE, '部门 → 工资 / 折旧计入', 'FF2E75B6')
    header(ws, 4, [(DP_NAME, '部门'), (DP_WAGE, '工资计入'), (DP_DEP, '折旧计入'), (DP_NOTE, '说明')], 'FF2E75B6', height=44)
    for i in range(DP_R1 - DP_R0 + 1):
        r = DP_R0 + i
        x = fixes.DEPTS[i] if i < len(fixes.DEPTS) else (None,) * 4
        put(ws, f'{DP_NAME}{r}', x[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{DP_WAGE}{r}', x[1], F_IN, FILL_IN, '@', AC)
        put(ws, f'{DP_DEP}{r}', x[2], F_IN, FILL_IN, '@', AC)
        put(ws, f'{DP_NOTE}{r}', x[3] or None, F_NOTE, FILL_IN, align=ALW)

    _block_title(ws, PT_NAME, PT_ACC, '人员类别 → 应付工资', 'FF2E75B6')
    header(ws, 4, [(PT_NAME, '人员类别'), (PT_ACC, '应付科目')], 'FF2E75B6', height=44)
    for i in range(PT_R1 - PT_R0 + 1):
        r = PT_R0 + i
        x = fixes.PERSON_TYPES[i] if i < len(fixes.PERSON_TYPES) else (None, None)
        put(ws, f'{PT_NAME}{r}', x[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{PT_ACC}{r}', x[1], F_IN, FILL_IN, '@', AC)

    _block_title(ws, FC_NAME, FC_NOTE, '固定资产类别', 'FF7F6000')
    header(ws, 4, [(FC_NAME, '类别'), (FC_ACC, '科目'), (FC_YEARS, '折旧\n年限'), (FC_SALV, '残值率'), (FC_NOTE, '说明')],
           'FF7F6000', height=44)
    for i in range(FC_R1 - FC_R0 + 1):
        r = FC_R0 + i
        x = fixes.FA_CATS[i] if i < len(fixes.FA_CATS) else (None,) * 5
        put(ws, f'{FC_NAME}{r}', x[0], F_IN, FILL_IN, align=AC)
        put(ws, f'{FC_ACC}{r}', x[1], F_IN, FILL_IN, '@', AC)
        put(ws, f'{FC_YEARS}{r}', x[2], F_IN, FILL_IN, '0', AC)
        put(ws, f'{FC_SALV}{r}', x[3], F_IN, FILL_IN, '0%', AC)
        put(ws, f'{FC_NOTE}{r}', x[4] or None, F_NOTE, FILL_IN, align=ALW)

    _block_title(ws, CFI_NO, CFI_SIDE, '现金流量项目（会小企03 行次）', C_RPT)
    header(ws, 4, [(CFI_NO, '行次'), (CFI_NAME, '项目'), (CFI_SIDE, '收/付')], C_RPT, height=44)
    for i in range(CFI_R1 - CFI_R0 + 1):
        r = CFI_R0 + i
        x = fixes.CF_ITEMS[i] if i < len(fixes.CF_ITEMS) else (None, None, None)
        put(ws, f'{CFI_NO}{r}', x[0], F_TXT, FILL_AUTO, '0', AC)
        put(ws, f'{CFI_NAME}{r}', x[1], F_TXT, FILL_AUTO, align=AL)
        put(ws, f'{CFI_SIDE}{r}', x[2], F_TXT, FILL_AUTO, align=AC)
    build_subs(ws, ctx)
    ws.auto_filter.ref = f'A4:{RC_NOTE}{RC_R1}'      # 用户 10/05 加的筛选（A075 照留）
    ws.freeze_panes = 'C5'
    return ws


def build_subs(ws, ctx):
    """10/15：「二级明细对照」块（【记账规则】AM:AQ）。数据＝ctx['subs']（ref_a075.SUBS）：[(类别, 二级, 科目, 说明)]"""
    _block_title(ws, SB_CAT, SB_KEY, '二级明细对照（流水「二级/备注」→ 明细科目）', 'FF2E75B6')
    header(ws, 4, [(SB_CAT, '收支类别'), (SB_SUB, '二级\n(流水 R 列写的)'), (SB_ACC, '记到科目'), (SB_NOTE, '说明'),
                   (SB_KEY, '键（自动）')], 'FF2E75B6', height=44)
    subs = ctx.get('subs') or []
    assert len(subs) <= SUB_R1 - SUB_R0 + 1
    for i in range(SUB_R1 - SUB_R0 + 1):
        r = SUB_R0 + i
        x = subs[i] if i < len(subs) else (None,) * 4
        put(ws, f'{SB_CAT}{r}', x[0], F_IN, FILL_IN, align=AL)
        put(ws, f'{SB_SUB}{r}', x[1], F_IN, FILL_IN, '@', AL)
        put(ws, f'{SB_ACC}{r}', x[2], F_IN, FILL_IN, '@', AC)
        put(ws, f'{SB_NOTE}{r}', x[3], F_NOTE, FILL_IN, align=ALW)
        put(ws, f'{SB_KEY}{r}', f'=IF(OR({SB_CAT}{r}="",{SB_SUB}{r}=""),"",TRIM({SB_CAT}{r})&"|"&TRIM({SB_SUB}{r}))',
            F_NOTE, FILL_AUTO, align=AL)
    add_list_dv(ws, f'{SB_CAT}{SUB_R0}:{SB_CAT}{SUB_R1}', '=类别列表')
    add_list_dv(ws, f'{SB_ACC}{SUB_R0}:{SB_ACC}{SUB_R1}', '=科目选择', '填编码或从下拉选', stop=False)
    k = f'${SB_KEY}${SUB_R0}:${SB_KEY}${SUB_R1}'
    ws.conditional_formatting.add(f'{SB_CAT}{SUB_R0}:{SB_SUB}{SUB_R1}',
                                  FormulaRule(formula=[f'AND(${SB_KEY}{SUB_R0}<>"",COUNTIF({k},${SB_KEY}{SUB_R0})>1)'], fill=FILL_WARN))
    ws.column_dimensions[SB_KEY].hidden = True


# ───────────────────────── 货柜登记（10/15，录入本：纯手填备查，不记账） ─────────────────────────
def build_container(wb, ctx):
    """【货柜登记】：签收日期、柜号、提单号、封条号、货物/供应商、备注，300 行，没有公式。数据＝ctx['containers']（导入/u2_货柜.json）。
    已经有这张表就不再建（p_a075.apply 里调；build 里另调也不会建两张）"""
    if SH_BOX in wb.sheetnames:
        return wb[SH_BOX]
    ws = wb.create_sheet(SH_BOX)
    widths(ws, {c: w for c, _, w in BOX_COLS})
    last = BOX_COLS[-1][0]
    title(ws, '货柜登记（货柜签收备查 · 只登记，不记账）', last, 'FF00B050',
          '💡 到一个货柜记一行：签收日期、柜号、提单号、封条号、什么货 / 谁家的。只是备查，不出凭证；货的钱在【采购入库】和资金台帐里记。'
          '往下面空行接着填，别插行删行。')
    header(ws, BOX_HDR, [(c, t) for c, t, _ in BOX_COLS], 'FF00B050', height=30)
    rows = ctx.get('containers') or []
    assert len(rows) <= BOX_R1 - BOX_R0 + 1
    keys = ['date', 'box', 'bl', 'seal', 'goods', 'note']
    for i in range(BOX_R1 - BOX_R0 + 1):
        r = BOX_R0 + i
        x = rows[i] if i < len(rows) else {}
        for (c, _, _), k in zip(BOX_COLS, keys):
            ws[f'{c}{r}'] = x.get(k)
    style_rows(ws, BOX_R0, BOX_R1, [c for c, _, _ in BOX_COLS], fmts={'A': DATE, 'B': '@', 'C': '@', 'D': '@'},
               aligns={'E': AL, 'F': AL})
    add_date_dv(ws, f'A{BOX_R0}:A{BOX_R1}')
    ws.auto_filter.ref = f'A{BOX_HDR}:{last}{BOX_R1}'
    ws.freeze_panes = f'A{BOX_R0}'
    return ws


# ───────────────────────── 月度汇率 ─────────────────────────
def build_rates(wb, ctx):
    ws = wb.create_sheet(SH_RATE)
    ninv = CI(RATE_INV0)
    widths(ws, {'A': 12, **{CL(i + 2): 12 for i in range(len(RATE_CURS))},
                **{CL(ninv + i): 12 for i in range(len(RATE_CURS))}, CL(ninv + len(RATE_CURS)): 45})
    title(ws, f'月度记账汇率（1 {BASE_NAME} = 多少单位外币）', 'I', C_BASE,
          f'💡 按日记账口径填写：1 {BASE_CUR} = X 原币。{BASE_CUR} 固定为 1，CNY / ZiG 等填实际报价；流水金额÷汇率得{BASE_NAME}，'
          f'月末调汇自动采用相应的折{BASE_NAME}系数。空白代表当月汇率待补，禁止臆测。')
    put(ws, 'A4', '月份', F_HDR, fill(C_BASE), align=AC)
    header(ws, RATE_HDR, [('A', '月份')] + [(CL(i + 2), c) for i, c in enumerate(RATE_CURS)], C_BASE, height=24)
    # A075：K..R＝1 单位外币折多少美元（＝1/汇率），记账汇率区指这一块；隐藏（A075 原件这几格表头沿用 K5 原来说明文字的灰色小字，数据格没设格式）
    for i, c in enumerate(RATE_CURS):
        if i == 0:
            put(ws, f'{CL(ninv + i)}{RATE_HDR}', f'{c}折{BASE_NAME}', F_NOTE, align=ALW, border=False)
        else:
            ws[f'{CL(ninv + i)}{RATE_HDR}'] = f'{c}折{BASE_NAME}'
    rates = ctx['rates']
    for n in range(1, N_MONTHS + 1):
        r = RATE_R0 + n - 1
        put(ws, f'A{r}', f'=EDATE(起始月份,{n - 1})', F_TXTB, FILL_SUBH, MONTH, AC)
        v = rates[n - 1] if n - 1 < len(rates) else {}
        for i, c in enumerate(RATE_CURS):
            put(ws, f'{CL(i + 2)}{r}', v.get(c), F_IN, FILL_IN if c != BASE_CUR else FILL_AUTO, RATE, AR_)
            src = f'{CL(i + 2)}{r}'
            ws[f'{CL(ninv + i)}{r}'] = f'=IF({src}="","",IF({src}>0,1/{src},""))'
    # A062 时 K5:K9 合并写说明；A075 改成表头后合并没拆，合并区里 K6:K9 照样有公式（照原件留着：
    # 不能用 merge_cells，那会把 K6:K9 的公式清掉；只登记合并区）。优化时可以拆掉。
    from openpyxl.worksheet.cell_range import CellRange
    ws.merged_cells.add(CellRange('K5:K9'))
    hide(ws, [CL(ninv + i) for i in range(len(RATE_CURS))])
    ws.freeze_panes = f'B{RATE_R0}'
    return ws


# ───────────────────────── 科目表 ─────────────────────────
def build_accounts(wb, ctx):
    ws = wb.create_sheet(SH_ACC)
    widths(ws, {AC_CODE: 11, AC_NAME: 26, AC_CLASS: 12, AC_DIR: 6, AC_AUX: 9, AC_FX: 12, AC_CASHF: 7, AC_STAT: 6,
                AC_PARA: 7, AC_LVL: 5, AC_FULL: 40, AC_LEAF: 5, AC_OPEN: 16, AC_KEY: 12, 'O': 34})
    title(ws, '会计科目表（原表2【2 科目表】287 个 ＋ 采矿业务补的明细）', AC_OPEN, C_BASE,
          '💡 编码按 4-2-2-2 分级。要加明细科目：往下面空行填（编码写成文本，比如 160110），不要插行删行。'
          '「期初余额」只在最末级科目填（借方正数、贷方负数），是建账起始月份那天的余额（现在都是 0，流水从 2025-05 起全量记账）。'
          '灰色几列自动：级次、全称、是不是末级。')
    header(ws, AC_HDR, [(AC_CODE, '编码'), (AC_NAME, '名称'), (AC_CLASS, '类别'), (AC_DIR, '余额\n方向'),
                        (AC_AUX, '辅助核算'), (AC_FX, '外币核算'), (AC_CASHF, '现金\n科目'), (AC_STAT, '状态'),
                        (AC_PARA, '平行\n科目'), (AC_LVL, '级次'), (AC_FULL, '全称'), (AC_LEAF, '末级'),
                        (AC_OPEN, '期初余额\n(借正贷负)'), (AC_KEY, '排序键'), ('O', '下拉用（编码 名称）')], C_BASE, height=36)
    accs = ctx['chart']
    A, B = acc(AC_CODE), acc(AC_NAME)
    for i in range(AC_R1 - AC_R0 + 1):
        r = AC_R0 + i
        a = accs[i] if i < len(accs) else {}
        c = put(ws, f'{AC_CODE}{r}', a.get('code'), F_IN, FILL_IN, '@', AL)
        put(ws, f'{AC_NAME}{r}', a.get('name'), F_IN, FILL_IN, align=AL)
        for col, key in ((AC_CLASS, 'cls'), (AC_DIR, 'dir'), (AC_AUX, 'aux'), (AC_FX, 'fx'), (AC_CASHF, 'cashf'),
                         (AC_STAT, 'stat'), (AC_PARA, 'para')):
            put(ws, f'{col}{r}', a.get(key), F_IN, FILL_IN, align=AC)
        x = f'{AC_CODE}{r}'
        put(ws, f'{AC_LVL}{r}', f'=IF({x}="","",(LEN({x})-2)/2)', F_AUTO, FILL_AUTO, '0', AC)
        parts = []
        for L in (4, 6, 8, 10):
            parts.append(f'IF(LEN({x})>={L},IF({L}>4,"-","")&IFERROR(INDEX({B},MATCH(LEFT({x},{L}),{A},0)),"?"),"")')
        put(ws, f'{AC_FULL}{r}', f'=IF({x}="","",' + '&'.join(parts) + ')', F_AUTO, FILL_AUTO, align=AL)
        put(ws, f'{AC_LEAF}{r}', f'=IF({x}="","",IF(COUNTIF({A},{x}&"?*")>0,"否","是"))', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{AC_OPEN}{r}', a.get('open'), F_IN, FILL_IN, MONEY, AR_)
        put(ws, f'{AC_KEY}{r}', f'=IF({x}="","",IFERROR(VALUE(LEFT({x}&"0000000000",10)),""))', F_NOTE, FILL_AUTO, '0', AC)
        # 隐藏 P：生产成本（4001/4101）末级科目的排序键——月末结转③按它自动找要转原矿的科目，新加的明细也会转
        put(ws, f'P{r}', f'=IF(AND(OR(LEFT({x},4)="4001",LEFT({x},4)="4101"),{AC_LEAF}{r}="是"),{AC_KEY}{r},"")', F_NOTE, FILL_AUTO, '0', AC)
        put(ws, f'O{r}', f'=IF({x}="","",{x}&" "&{AC_NAME}{r})', F_NOTE, FILL_AUTO, align=AL)
        # 隐藏 Q（opt）：这一行在【科目余额表】里有没有自己的一行（排序键有、而且是这个键第一次出现）——
        # 记账分录「余额键」按它分：有＝按科目行号汇总；没有（编码格式不对、补零撞键）＝按编码前缀汇总（跟原来的通配符一样）
        put(ws, f'Q{r}', f'=IF({AC_KEY}{r}="",0,IF(MATCH({AC_KEY}{r},${AC_KEY}${AC_R0}:${AC_KEY}${AC_R1},0)={r - AC_R0 + 1},1,0))',
            F_NOTE, FILL_AUTO, '0', AC)
        if a.get('added'):
            ws[f'{AC_NAME}{r}'].font = Font(name=YH, sz=10, color='FF1F4E79', bold=True)
    ws.conditional_formatting.add(f'{AC_CODE}{AC_R0}:{AC_CODE}{AC_R1}',
                                  FormulaRule(formula=[f'AND({AC_CODE}{AC_R0}<>"",COUNTIF({A},{AC_CODE}{AC_R0})>1)'],
                                              fill=FILL_WARN))
    add_list_dv(ws, f'{AC_DIR}{AC_R0}:{AC_DIR}{AC_R1}', '"借,贷"')
    ws.column_dimensions['O'].hidden = True
    ws.column_dimensions['P'].hidden = True
    put(ws, f'Q{AC_HDR}', '余额表行', F_NOTE, FILL_AUTO, align=AC)
    ws.column_dimensions['Q'].hidden = True
    ws.auto_filter.ref = f'A{AC_HDR}:{AC_KEY}{AC_R1}'
    ws.freeze_panes = f'C{AC_R0}'
    return ws


# ───────────────────────── 定义名称 ─────────────────────────
def define_names(wb):
    def dn(name, ref):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)

    for key, _, nm in PARAMS:
        dn(nm, P(key))
    dn('报表年度', HOME_YEAR)
    for i, sh in enumerate(USED):                       # opt：各录入表末行（_辅助 第 206 行起）
        dn(USED[sh][0], f"{q(SH_AUX)}!$B${USED_R0 + i}")
    dn('最后有数月序', f"{q(SH_AUX)}!$B${USED_R0 + len(USED)}")
    dn('流水块前原币', f"{q(SH_AUX)}!$C${CB_R0}:${CL(2 + N_ACC)}${CB_R0 + N_CB - 1}")
    dn('流水块前美元', f"{q(SH_AUX)}!${CL(23 + 1)}${CB_R0}:${CL(23 + N_ACC)}${CB_R0 + N_CB - 1}")
    dn('报表月份', HOME_MONTH)
    b = lambda col, r0, r1: rng(SH_BASE, col, r0, r1)
    rr = lambda col, r0, r1: rng(SH_RULE, col, r0, r1)
    # 跟原表【1现金流水总表】公式里用的名字一样（原来指向另一个文件，现在指向本册）
    dn('账户名', b(A_NAME, ACC_R0, ACC_R1))
    dn('账户银行', b(A_BANK, ACC_R0, ACC_R1))
    dn('账户币种', b(A_CUR, ACC_R0, ACC_R1))
    dn('账户期初', b(A_OPEN, ACC_R0, ACC_R1))
    dn('账户期初本位', b(A_OPENB, ACC_R0, ACC_R1))
    dn('账户科目', b(A_CODE, ACC_R0, ACC_R1))
    dn('类别表', rr(RC_NAME, RC_R0, RC_R1))
    dn('类别属性', rr(RC_ATTR, RC_R0, RC_R1))
    dn('项目表', b(PJ_NAME, PRJ_R0, PRJ_R1))
    dn('单位表', b(CP_NAME, CP_R0, CP_R1))
    dn('汇率月份', rng(SH_RATE, 'A', RATE_R0, RATE_R1))
    dn('汇率币种头', f"{q(SH_RATE)}!$B${RATE_HDR}:${CL(1 + len(RATE_CURS))}${RATE_HDR}")
    # A075：记账汇率区＝「1 单位外币折多少美元」那一块（K:R），月末调汇、资金账户余额表按它乘
    dn('记账汇率区', f"{q(SH_RATE)}!${RATE_INV0}${RATE_R0}:${CL(CI(RATE_INV0) + len(RATE_CURS) - 1)}${RATE_R1}")

    def lst(sh, col, r0, r1):
        # A075：下拉名称从 OFFSET 改成「首格:INDEX(…)」写法（不易失；Excel 2010 / WPS 都认）
        rg = f'{q(sh)}!${col}${r0}:${col}${r1}'
        return f'{q(sh)}!${col}${r0}:INDEX({rg},MAX(1,IFERROR(LOOKUP(2,1/({rg}<>""),ROW({rg}))-{r0 - 1},1)))'
    dn('账户列表', lst(SH_BASE, A_NAME, ACC_R0, ACC_R1))
    dn('类别列表', lst(SH_RULE, RC_NAME, RC_R0, RC_R1))
    dn('项目列表', lst(SH_BASE, PJ_NAME, PRJ_R0, PRJ_R1))
    dn('往来列表', lst(SH_BASE, CP_NAME, CP_R0, CP_R1))
    dn('物料列表', lst(SH_BASE, M_NAME, MAT_R0, MAT_R1))
    dn('物料类别列表', lst(SH_RULE, MC_NAME, MC_R0, MC_R1))
    dn('用途列表', lst(SH_RULE, US_NAME, US_R0, US_R1))
    dn('部门列表', lst(SH_RULE, DP_NAME, DP_R0, DP_R1))
    dn('人员类别列表', lst(SH_RULE, PT_NAME, PT_R0, PT_R1))
    dn('资产类别列表', lst(SH_RULE, FC_NAME, FC_R0, FC_R1))
    dn('科目选择', lst(SH_ACC, 'O', AC_R0, AC_R1))
    # 10/15：二级明细对照（s_cash.raw_formula 用：INDEX(二级_科目,MATCH(E&"|"&R,二级_键,0))）
    dn('二级_键', rr(SB_KEY, SUB_R0, SUB_R1))
    dn('二级_科目', rr(SB_ACC, SUB_R0, SUB_R1))
