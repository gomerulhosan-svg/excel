# -*- coding: utf-8 -*-
"""拆本：录入本（A075-1_录入本.xlsx）的写法 ＋ 录入本专用表（首页、【录入检查】）。

录入本＝一本版删掉报表表，录入表里只留「轻」公式（带出银行、币种、汇率、余额、核对…），只为记账 / 报表算的「重」列清空。
这里管三件事：
  1. 录入本里要换写法的公式（ENTRY_FORMULAS）：现金流水 N/O 余额和 _辅助 块前累计按行号算块、各「同一单据号 / 合并号第一行」
     按行号找（按日期整行排序以后照样对）；领用 O、销售 Q、流水 N3 不再依赖只在报表本算的列。
  2. 录入本里清空的列（ENTRY_CLEAR）：重列（照旧隐藏）＋只在报表本算的列（ENTRY_REPORT_ONLY：灰格，表头加「（在报表本看）」）。
  3. 录入本首页（build_entry_home）、【录入检查】（build_entry_check）。

──────────── 接口（给 split / biz）────────────
  apply(wb, ctx=None) → [日志]      一次做完 1、2，再补名称。录入本 wb 在 build 内存里，顺序：
                                     build() 全部做完（含 p1008 / p_a075 / p_opt / p_wps.apply）→ 删报表表 → entry.apply(wb)
                                     → build_entry_home(wb, ctx) → build_entry_check(wb, ctx) → 过滤名称（只留用到的）→ 存盘。
                                     （p_wps.apply 在 apply 之后跑也行：这里写的公式里没有「表名!」和往下撑的区域，用到的名称两边都认。）
  ENTRY_FORMULAS {(表名, 列): fn(行号, 原公式=None) → 公式}
                                     行号范围见 ENTRY_ROWS。给了原公式（一本版 build 完的那一格）就在它上面换掉要换的那一段
                                     （biz 在 p_a075 里改过的提示字跟着留下）；不给就从 s_cash / s_mods 重新生成。
  ENTRY_ROWS     {(表名, 列): [行号…]}
  ENTRY_CLEAR    {表名: [列…]}       录入本清空（数据行，表头留着）
  ENTRY_REPORT_ONLY {表名: [列…]}    其中只在报表本算的（领用 I/J、销售 M/N、固定资产 M/N/O）
  AUX_KEEP       _辅助 留下的行（末行 206～214、块前累计 218～259），别的格子 apply 清掉
  ensure_names(wb)                   补这里用到的名称（分段名称、「区_表_A1_B2」），apply / build_entry_* 里已经调过
  build_entry_home(wb, ctx) → ws     录入本【首页】（有旧的首页就换掉，放第一张）
  build_entry_check(wb, ctx) → ws    录入本【录入检查】（有就换掉；放在【使用说明】前面，没有就放最后）
  ECHK、ECHK_ROWS、ECHK_MIRROR       【录入检查】表名、各项在第几行、报表本取数表要取的格子：
                                     报表本里有一张隐藏取数表【录入检查】，同名同位置，取 ECHK_MIRROR 这些格子
                                     （写法跟别的取数格一样：=IF('[1]录入检查'!C5="","",'[1]录入检查'!C5)）；
                                     s_check.report_check41(wb) 读 C3、C{ECHK_ROWS['hdr']}、C{ECHK_ROWS['rows']}。
报表本这边（s_check.py）：report_check41(wb) 改数据校验第 41 项；home_link_line(ws, row) 在报表本首页写「数据来自录入本：流水 N 笔，最后一笔 yyyy/m/d」。
"""
import re

from openpyxl.workbook.defined_name import DefinedName
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule

from common import *
import s_cash
import s_mods
import s_check
from p_wps import SHEET_REF, CELL, _unq, _safe, _absref

ECHK = '录入检查'
GREY = fill('FFD9D9D9')            # 录入本里「在报表本看」的灰格
REPORT_NOTE = '（在报表本看）'

# ─────────────────────────── 名称 ───────────────────────────
ENTRY_AREAS = {}                   # 这里新建的「区_表_A1_B2」名称 → 引用（跟 p_wps 的命名一样，同一区域两边是同一个名字）


def _area(sh, a, b=None):
    nm = f'区_{_safe(sh)}_{a.replace("$", "")}' + (f'_{b.replace("$", "")}' if b else '')
    ref = f"{q(sh)}!{_absref(a)}" + (f':{_absref(b)}' if b else '')
    old = ENTRY_AREAS.setdefault(nm, ref)
    assert old == ref, (nm, old, ref)
    return nm


def _wps(f, title):
    """公式里的「别的表!$A$5:$A$454」（绝对行）换成名称，本表名去掉；跨表相对引用不该出现（报错）。成列的公式 WPS 才不拆"""
    parts = f.split('"')
    for i in range(0, len(parts), 2):
        def sub(m):
            sh, a, b = _unq(m.group('sh')), m.group('a'), m.group('b')
            if sh == title:
                return a + (f':{b}' if b else '')
            ma, mb = CELL.match(a), (CELL.match(b) if b else None)
            if ma.group(3) == '$' and (mb is None or mb.group(3) == '$'):
                return _area(sh, a, b)
            raise AssertionError(f'录入本：{title} 的公式里有跨表相对引用 {m.group(0)}：{f[:160]}')
        parts[i] = SHEET_REF.sub(sub, parts[i])
    return '"'.join(parts)


def ensure_names(wb):
    """补名称：分段名称（common.SEG_NAMES，p_wps.apply 已经写过的才补——没跑过的让 p_wps 自己写）＋ 这里建的「区_」名称"""
    n = 0
    if any(nm in wb.defined_names for nm in list(SEG_NAMES)[:20]):
        for nm, ref in SEG_NAMES.items():
            if nm not in wb.defined_names:
                wb.defined_names[nm] = DefinedName(nm, attr_text=ref, hidden=True)
                n += 1
    for nm, ref in ENTRY_AREAS.items():
        if nm not in wb.defined_names:
            wb.defined_names[nm] = DefinedName(nm, attr_text=ref, hidden=True)
            n += 1
    return n


def F(sh, col):
    """录入表某一整列的名称（流水_L＝现金流水总表!$L$6:$L$3965）"""
    return full_name(sh, col)


# ─────────────────────────── 录入本写法 ───────────────────────────
def _head_e(sh, dcol, kcol, r, dollar=''):
    """同一个单据号 / 合并号的第一行的日期——按行号找：MATCH(键,第一行到本行,0)（排序、插行以后照样对；键空着不找）"""
    k = f'{dollar}{kcol}{r}'
    return f'IF({k}="","",INDEX({F(sh, dcol)},MATCH({k},INDEX({F(sh, kcol)},1):${kcol}{r},0)))'


_BLK = f'(INT((ROW()-{CASH_R0})/{CB_N})+1)'                 # 本行在第几块（按行号算，不写死）
_BLK0 = f'(INT((ROW()-{CASH_R0})/{CB_N})*{CB_N}+1)'         # 这块第一行是第几笔


def _blk_loc(col, r):
    return f'INDEX({F(SH_CASH, col)},{_BLK0}):${col}{r}'


def cash_n(r, old=None):
    """N 账户余额(原币)：期初＋块前累计（_辅助，块号按行号算）＋本块第一行加到本行；账户不在档案里的从第一行加到本行（跟一本版一样）"""
    if r == 3:          # N3「生成凭证的笔数」＝有金额的笔数 − 其中核对是 ✗ 的（一本版是 SUM(有效 AC)，AC 只在报表本算）
        m, s = cash_u(K_UZS), cash_u(K_CHK)
        return (f'=COUNTIF({m},">0")+COUNTIF({m},"<0")-COUNTIFS({m},">0",{s},"✗*")-COUNTIFS({m},"<0",{s},"✗*")')
    mt = f'MATCH($B{r},账户名,0)'
    allup = f'SUMIFS(INDEX({F(SH_CASH, "L")},1):$L{r},INDEX({F(SH_CASH, "B")},1):$B{r},$B{r})'
    return (f'=IF($B{r}="","",ROUND(SUMIF(账户名,$B{r},账户期初)+IFERROR(INDEX(流水块前原币,{_BLK},{mt})'
            f'+SUMIFS({_blk_loc("L", r)},{_blk_loc("B", r)},$B{r}),{allup}),2))')


def cash_o(r, old=None):
    """O 账面美元余额：同 N"""
    mt = f'MATCH($B{r},账户名,0)'
    return (f'=IF($B{r}="","",IFERROR(ROUND(INDEX(账户期初本位,{mt})+INDEX(流水块前美元,{_BLK},{mt})'
            f'+SUMIFS({_blk_loc("M", r)},{_blk_loc("B", r)},$B{r}),2),""))')


def _swap(old, a, b, what):
    assert a in old, f'录入本：{what} 原公式里没找到要换的那一段：{a[:80]} … 原公式：{old[:160]}'
    return old.replace(a, b)


def cash_s(r, old=None):
    """S 核对：只把「同一合并号第一笔的日期」换成按行号找"""
    new = _head_e(SH_CASH, 'A', 'V', r, '$')
    if old is None:
        return _wps(s_cash.chk_formula(r, new), SH_CASH)
    return _swap(old, first_val(SH_CASH, f'$V{r}', 'V', 'A', r), new, f'{SH_CASH}!S{r}')


def aux_blk(col):
    """_辅助 块前累计（第 220 行起）：第 b 块以前各账户的原币 / 折美元合计；块按「流水第几笔」（INDEX(整列名称,第几笔)）算，不用段名称——
    插行、删行以后名称伸缩，块的边界跟流水 N/O 的「按行号算块」对得上"""
    ci = CI(col)
    src, a = ('L', ci - 2) if ci <= 2 + N_ACC else ('M', ci - 23)
    assert 1 <= a <= N_ACC

    def fn(r, old=None):
        if r == CB_R0:
            return 0
        an = f'INDEX(账户名,COLUMN()-{ci - a})'
        p0, p1 = f'(ROW()-{CB_R0 + 1})*{CB_N}+1', f'(ROW()-{CB_R0})*{CB_N}'
        rg = lambda c: f'INDEX({F(SH_CASH, c)},{p0}):INDEX({F(SH_CASH, c)},{p1})'
        return f'=IF({an}="",0,{col}{r - 1}+SUMIFS({rg(src)},{rg("B")},{an}))'
    return fn


def _month_end1(d):
    return f'DATE(YEAR({d}),MONTH({d})+1,1)'


def iss_o(r, old=None):
    """领用 O 校验：同一领料单号第一行按行号找；「单价为 0」（I 列只在报表本算）改成「期初是 0、到本月底采购入库里也没进过这个物料」"""
    head = _head_e(SH_ISS, I_DATE, I_NO, r)
    cond = (f'AND(N(INDEX({_area(SH_BASE, f"${M_Q0}${MAT_R0}", f"${M_Q0}${MAT_R1}")},'
            f'MATCH({I_MAT}{r},{_area(SH_BASE, f"${M_NAME}${MAT_R0}", f"${M_NAME}${MAT_R1}")},0)))=0,'
            f'COUNTIFS({mod_u(SH_BUY, B_MAT)},{I_MAT}{r},{mod_u(SH_BUY, B_DATE)},"<"&{_month_end1(f"{I_DATE}{r}")})=0)')
    msg = '⚠ 库里还没有这个物料（期初是 0，到这个月底也没录过采购入库），先录采购入库'
    if old is None:
        return _wps(s_mods.iss_chk(r, head, zero=(cond, msg)), SH_ISS)
    new = _swap(old, s_mods._head(SH_ISS, I_DATE, I_NO, r), head, f'{SH_ISS}!O{r}')
    new, n = re.subn(rf'IF\(N\({I_PRICE}{r}\)=0,"[^"]*",', lambda m: f'IF({cond},"{msg}",', new)
    assert n == 1, f'录入本：{SH_ISS}!O{r} 没找到「单价为 0」那一条：{old[:160]}'
    return new


def sale_q(r, old=None):
    """销售 Q 校验：同一结算单号第一行按行号找；「没有成本」（M 列只在报表本算）改成「到本月还没登记过冶炼量」"""
    head = _head_e(SH_SALE, S_DATE, S_NO, r)
    cond = (f'COUNTIFS({mod_u(SH_PROD, PD_TON)},">0",{mod_u(SH_PROD, PD_DATE)},"<"&{_month_end1(f"{S_DATE}{r}")})=0')
    msg = f'⚠ 到这个月还没登记过冶炼量（【{SH_PROD}】），销售成本按 0'
    if old is None:
        return _wps(s_mods.sale_chk(r, head, nocost=(cond, msg)), SH_SALE)
    new = _swap(old, s_mods._head(SH_SALE, S_DATE, S_NO, r), head, f'{SH_SALE}!Q{r}')
    new, n = re.subn(rf'IF\(N\({S_UC}{r}\)=0,"[^"]*",', lambda m: f'IF({cond},"{msg}",', new)
    assert n == 1, f'录入本：{SH_SALE}!Q{r} 没找到「没有成本」那一条：{old[:160]}'
    return new


def buy_u(r, old=None):
    """采购 U 校验：同一单据号第一行按行号找"""
    head = _head_e(SH_BUY, B_DATE, B_NO, r)
    if old is None:
        return _wps(s_mods.buy_chk(r, head), SH_BUY)
    return _swap(old, s_mods._head(SH_BUY, B_DATE, B_NO, r), head, f'{SH_BUY}!U{r}')


_CASH_ROWS = list(range(CASH_R0, CASH_R1 + 1))
_AUX_COLS = [CL(2 + a) for a in range(1, N_ACC + 1)] + [CL(23 + a) for a in range(1, N_ACC + 1)]
ENTRY_FORMULAS = {
    (SH_CASH, K_BAL): cash_n,
    (SH_CASH, K_BBAL): cash_o,
    (SH_CASH, K_CHK): cash_s,
    (SH_BUY, B_CHK): buy_u,
    (SH_ISS, I_CHK): iss_o,
    (SH_SALE, S_CHK): sale_q,
    **{(SH_AUX, c): aux_blk(c) for c in _AUX_COLS},
}
ENTRY_ROWS = {
    (SH_CASH, K_BAL): [3] + _CASH_ROWS,
    (SH_CASH, K_BBAL): _CASH_ROWS,
    (SH_CASH, K_CHK): _CASH_ROWS,
    (SH_BUY, B_CHK): list(range(MOD_R0, BUY_R1 + 1)),
    (SH_ISS, I_CHK): list(range(MOD_R0, ISS_R1 + 1)),
    (SH_SALE, S_CHK): list(range(MOD_R0, SALE_R1 + 1)),
    **{(SH_AUX, c): list(range(CB_R0 + 1, CB_R0 + N_CB)) for c in _AUX_COLS},
}
ENTRY_CLEAR = {
    SH_CASH: [K_DR, K_CR, K_VALID, K_CFI, K_GSEQ, K_GHEAD, K_VID, K_VMEMO, K_MKEY, K_GCNT, K_GSUM, K_GMEMO, s_cash.K_SLOT],
    SH_BUY: [B_VALID, B_GKEY, B_LEGS, B_GPRE, B_VID, SEG[SH_BUY][3]],
    SH_ISS: [I_PRICE, I_AMT, I_VALID, I_GKEY, I_LEGS, I_GPRE, I_VID, SEG[SH_ISS][3]],
    SH_SALE: [S_UC, S_COGS, S_VALID, S_GKEY, S_LEGS, S_GPRE, S_VID, SEG[SH_SALE][3]],
    SH_PAY: [W_VALID, W_GKEY, W_LEGS, W_GPRE, W_VID, SEG[SH_PAY][3]],
    SH_FA: [F_MONTHS, F_ACCUM, F_NBV, F_END, F_LASTAMT, F_OK],
    SH_MAN: [MN_SEQ, MN_OK, MN_VID, MN_KEY, MN_VMEMO, SEG[SH_MAN][3]],
}
ENTRY_REPORT_ONLY = {SH_ISS: [I_PRICE, I_AMT], SH_SALE: [S_UC, S_COGS], SH_FA: [F_MONTHS, F_ACCUM, F_NBV]}
AUX_KEEP = (range(USED_R0, USED_R0 + len(USED) + 1), range(CB_R0 - 2, CB_R0 + N_CB))
_R0 = {SH_CASH: CASH_R0}


def apply(wb, ctx=None):
    log = []
    n = 0
    for (sh, col), fn in ENTRY_FORMULAS.items():
        ws = wb[sh]
        for r in ENTRY_ROWS[(sh, col)]:
            c = ws[f'{col}{r}']
            old = c.value if isinstance(c.value, str) and c.value.startswith('=') else None
            c.value = fn(r, old)
            n += 1
    log.append(f'录入本写法 {n:,} 格（流水 N/O 余额、_辅助 块前累计按行号算块；流水 S、采购 U、领用 O、销售 Q 的「同号第一行」按行号找；'
               f'领用 O、销售 Q、流水 N3 不用报表本才算的列）')
    n = 0
    for sh, cols in ENTRY_CLEAR.items():
        ws = wb[sh]
        r0 = _R0.get(sh, MOD_R0)
        for col in cols:
            for r in range(r0, ws.max_row + 1):
                c = ws[f'{col}{r}']
                if c.value is not None:
                    c.value = None
                    n += 1
    for sh, cols in ENTRY_REPORT_ONLY.items():
        ws = wb[sh]
        r0 = _R0.get(sh, MOD_R0)
        for col in cols:
            h = ws[f'{col}{r0 - 1}']
            if REPORT_NOTE not in str(h.value or ''):
                h.value = f'{h.value or ""}\n{REPORT_NOTE}'
            for r in range(r0, mod_r1(sh) + 1):
                ws[f'{col}{r}'].fill = GREY
    log.append(f'清空 {n:,} 格（重列照旧隐藏；领用 I/J、销售 M/N、固定资产 M/N/O 灰格，表头加「{REPORT_NOTE}」）')
    ws = wb[SH_AUX]
    keep = set().union(*AUX_KEEP)
    n = 0
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None and c.row not in keep:
                c.value = None
                n += 1
    log.append(f'_辅助 只留末行（{USED_R0}～{USED_R0 + len(USED)}）和块前累计（{CB_R0 - 2}～{CB_R0 + N_CB - 1}），清掉 {n:,} 格')
    log.append(f'补名称 {ensure_names(wb)} 个')
    return log


# ─────────────────────────── 录入检查 ───────────────────────────
def _probe_cols(wb):
    """各表「表头那一行最右边的录入列」：在它左边插删列、在表头上面插删行，它就挪了"""
    rule_last = CFI_SIDE
    for nm in ('二级_科目', '二级_键'):        # biz 加的「二级明细对照」在记账规则最右边就用它
        dn = wb.defined_names.get(nm) if hasattr(wb.defined_names, 'get') else None
        m = re.search(r'!\$?([A-Z]+)\$?\d+', (dn.attr_text if dn else '') or '')
        if m and CI(m.group(1)) > CI(rule_last):
            rule_last = m.group(1)
    return [  # (表, 表头行, 探针列, 数据第一行, 数据最后一行)
        (SH_CASH, CASH_HDR, K_ATT, CASH_R0, CASH_R1), (SH_BUY, MOD_HDR, B_NOTE, MOD_R0, BUY_R1),
        (SH_SALE, MOD_HDR, S_NOTE, MOD_R0, SALE_R1), (SH_ISS, MOD_HDR, I_NOTE, MOD_R0, ISS_R1),
        (SH_PROD, MOD_HDR, PD_NOTE, MOD_R0, PROD_R1), (SH_PAY, MOD_HDR, W_NOTE, MOD_R0, PAY_R1),
        (SH_FA, MOD_HDR, F_REMK, MOD_R0, FA_R1), (SH_MAN, MOD_HDR, MN_ATT, MOD_R0, MAN_R1),
        (SH_BASE, 4, M_NOTE, 5, max(ACC_R1, PRJ_R1, CP_R1, MAT_R1)),
        (SH_RULE, 4, rule_last, 5, max(RC_R1, MC_R1, US_R1, DP_R1, PT_R1, FC_R1, CFI_R1)),
        (SH_RATE, RATE_HDR, CL(1 + len(RATE_CURS)), RATE_R0, RATE_R1), (SH_ACC, AC_HDR, AC_OPEN, AC_R0, AC_R1)]


ECHK_R0 = 5


def _echk_items(wb):
    cnt = lambda rg, p='✗': f'COUNTIF({rg},"{p}*")'
    tabs = [(SH_BUY, B_CHK, '购进'), (SH_SALE, S_CHK, '销售'), (SH_ISS, I_CHK, ''), (SH_PROD, PD_CHK, ''),
            (SH_PAY, W_CHK, ''), (SH_FA, F_CHK, ''), (SH_MAN, MN_CHK, '')]
    items = [
        ('cash_x', f'资金台帐【{SH_CASH}】：有 ✗ 的行', s_check._ok0(cnt(cash_u(K_CHK)), '✗ '),
         f'到【{SH_CASH}】最右边「核对」列筛选 ✗，按提示改（✗ 的行在报表本里不出凭证）'),
        ('cash_w', f'资金台帐【{SH_CASH}】：有 ⚠ 的行', s_check._ok0(cnt(cash_u(K_CHK), '⚠'), '⚠ '),
         '提醒，不一定是错：余额变负数、科目不是末级、源数据【待补】之类，看一眼'),
    ]
    for sh, col, alias in tabs:
        items.append((sh, f'{alias + "【" + sh + "】" if alias else "【" + sh + "】"}：有 ✗ 的行',
                      s_check._ok0(cnt(mod_u(sh, col)), '✗ '), f'到【{sh}】「校验」列筛选 ✗，按提示改'))
    # 最近一笔流水（日期最晚的）那个月：账户用到的外币汇率填了没有（原数据校验第 18 项「报表月汇率」，录入本没有报表月）
    mx = f'MAX({cash_u(K_DATE)})'
    nrow = '$E${r}'           # 隐藏辅助格：那个月是第几个月（{r} 写的时候换成本行）
    miss = '+'.join(f'IF(COUNTIF(账户币种,"{c}")>0,IF(N(INDEX(记账汇率区,{nrow},{i + 1}))=0,1,0),0)'
                    for i, c in enumerate(RATE_CURS) if c != BASE_CUR)
    ym = f'YEAR({mx})&"年"&MONTH({mx})&"月"'
    items.append(('rate', '最近一笔流水（日期最晚的）那个月的外币汇率填了没有',
                  f'=IF(COUNT({cash_u(K_DATE)})=0,"√ 还没有流水",IF({nrow}=0,"✗ 流水日期早于建账起始月份",'
                  f'IF(({miss})=0,"√ "&{ym}&"的汇率都填了","✗ "&{ym}&"有外币的汇率还没填")))',
                  f'到【{SH_RATE}】把这个月人民币、ZiG 的汇率填上（1 美元＝多少）；不填，这个月的流水折不了美元（核对列报 ✗）',
                  f'=IFERROR(MATCH({mx},汇率月份,1),0)'))
    for i, (name, f, how) in enumerate(s_check.base_items()):
        items.append((f'base{i + 1}', name, f, how))
    items.append(('cap', f'容量：【{SH_CASH}】剩余空行',
                  f'=IF({N_CASH}-{s_check._nfill(cash(K_DATE))}<200,"⚠ 只剩 "&({N_CASH}-{s_check._nfill(cash(K_DATE))})&" 行",'
                  f'"√ 还剩 "&({N_CASH}-{s_check._nfill(cash(K_DATE))})&" 行")', f'共 {N_CASH:,} 行，快满了找我加'))
    pr = _probe_cols(wb)
    hdr = '&'.join(f'IF(IFERROR(AND(ROW({q(sh)}!${col}${h})={h},COLUMN({q(sh)}!${col}${h})={CI(col)}),FALSE),"","【{sh}】")'
                   for sh, h, col, _, _ in pr)
    items.append(('hdr', '各表表头、列还在原位（表头上面没插删行，没插删列）',
                  '=IF($E${r}="","√ 都在原位","✗ "&$E${r}&"表头上面插删了行，或者插删了列")',
                  '报表本是按位置从这本取数的：表头上面插删了行、或者插删了列，报表本会读错位置。按 Ctrl+Z 撤销；'
                  '已经存盘了，就把多出来的行 / 列删掉、删掉的补回来，让表头回到原来那一行、原来那一列',
                  '=' + hdr))
    rows = '&'.join(f'IF(IFERROR(AND(ROW({q(sh)}!$A${r0})={r0},ROWS({q(sh)}!$A${r0}:$A${r1})={r1 - r0 + 1}),FALSE),"","【{sh}】")'
                    for sh, _, _, r0, r1 in pr)
    items.append(('rows', '数据区没插入 / 删除整行',
                  '=IF($E${r}="","√ 没有","⚠ "&$E${r}&"数据区插入或删除了整行")',
                  '报表本照样能读（按位置读，等于这几行换了位置）。但插进来的那一行，灰格子里没有公式：那一行看不到余额、核对，'
                  '余额也不算它。漏记的直接记在最后一行下面、日期写对就行，不用插行；要删一笔，选中那一行的内容按 Delete 清空',
                  '=' + rows))
    return items


def _echk_layout(wb=None):
    """各项在第几行（不建表也能算：报表本 s_check 用）"""
    from openpyxl import Workbook
    keys = [it[0] for it in _echk_items(wb or Workbook())]
    rows = {k: ECHK_R0 + i for i, k in enumerate(keys)}
    return rows, ['C3', f'C{ECHK_R0}:C{ECHK_R0 + len(keys) - 1}']


# ECHK_ROWS：键 → 行号（cash_x、cash_w、各录入表、rate、base1～5、cap、hdr、rows）；ECHK_MIRROR：报表本取数表【录入检查】要取的格子
ECHK_ROWS, ECHK_MIRROR = _echk_layout()


def build_entry_check(wb, ctx=None):
    if ECHK in wb.sheetnames:
        del wb[ECHK]
    idx = wb.sheetnames.index(SH_HELP) if SH_HELP in wb.sheetnames else len(wb.sheetnames)
    ws = wb.create_sheet(ECHK, idx)
    widths(ws, {'A': 5, 'B': 40, 'C': 46, 'D': 70, 'E': 12})
    title(ws, '录入检查（录完看一眼：✗ 要改，⚠ 看一眼）', 'D', C_CHK,
          '💡 这一本只管录。这里查的是录入的东西对不对、表格有没有被挪动；凭证、账簿、报表的检查在【报表本】的【数据校验】。'
          '报表本是按位置从这本取数的，所以别在表头上面插删行、别插删列。')
    ws.row_dimensions[3].height = 26
    for c in 'ABCD':
        put(ws, f'{c}3', None, fill_=FILL_KPI, border=False)
    header(ws, 4, [('A', '序'), ('B', '检查项目'), ('C', '结果'), ('D', '怎么处理')], C_CHK, height=24)
    items = _echk_items(wb)
    rows = {}
    for i, it in enumerate(items):
        key, name, f, how = it[:4]
        r = ECHK_R0 + i
        rows[key] = r
        put(ws, f'A{r}', i + 1, F_TXT, align=AC)
        put(ws, f'B{r}', name, F_TXT, align=AL)
        put(ws, f'C{r}', f.replace('{r}', str(r)), F_TXTB, align=AL)
        put(ws, f'D{r}', how, F_NOTE, align=ALW)
        if len(it) > 4:
            put(ws, f'E{r}', it[4].replace('{r}', str(r)), F_NOTE, border=False)
        ws.row_dimensions[r].height = 30 if len(how) < 70 else 44
    R1 = ECHK_R0 + len(items) - 1
    assert rows == ECHK_ROWS, (rows, ECHK_ROWS)
    ws.conditional_formatting.add(f'C{ECHK_R0}:C{R1}', FormulaRule(formula=[f'LEFT($C{ECHK_R0},1)="✗"'], fill=FILL_WARN,
                                                                    font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'C{ECHK_R0}:C{R1}', FormulaRule(formula=[f'LEFT($C{ECHK_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'C{ECHK_R0}:C{R1}', FormulaRule(formula=[f'LEFT($C{ECHK_R0},1)="√"'], fill=FILL_OK))
    put(ws, 'A3', '汇总：', F_KPI_L, FILL_KPI, align=AR_, border=False)
    ws.merge_cells('A3:B3')
    put(ws, 'C3', f'=IF(COUNTIF(C{ECHK_R0}:C{R1},"✗*")=0,"√ 没有 ✗","✗ "&COUNTIF(C{ECHK_R0}:C{R1},"✗*")&" 项要改")'
                  f'&"　　⚠ "&COUNTIF(C{ECHK_R0}:C{R1},"⚠*")&" 项提醒"', Font(name=YH, sz=12, bold=True, color='FFC00000'),
        FILL_KPI, align=AL, border=False)
    put(ws, f'E{ECHK_R0 - 1}', '（辅助）', F_NOTE, border=False)
    ws.column_dimensions['E'].hidden = True
    ws.freeze_panes = 'A5'
    ws.sheet_view.showGridLines = False
    ensure_names(wb)
    return ws


# ─────────────────────────── 录入本首页 ───────────────────────────
HOME_TABS = [  # (表, 做什么, 有没有「校验」列)
    (SH_CASH, '资金台帐：每一笔收钱、付钱记一行（美元、人民币、FBC 三个账户都记在这张）', (K_NET, K_CHK)),
    (SH_BUY, '购进：进一次货、一个物资记一行', (B_SEQ, B_CHK)),
    (SH_SALE, '销售：卖一次黄金记一行', (S_SEQ, S_CHK)),
    (SH_ISS, '领用：从库里领物资，领一次记一行', (I_SEQ, I_CHK)),
    (SH_PROD, '冶炼登记：每次冶炼出多少克金', (PD_SEQ, PD_CHK)),
    (SH_PAY, '工资：每月按部门记一次应发工资', (W_SEQ, W_CHK)),
    (SH_FA, f'设备卡片（折旧、净值{REPORT_NOTE}）', (None, F_CHK)),
    (SH_MAN, '调账用的凭证（一般用不着）', (MN_HAS, MN_CHK)),
    ('货柜登记', '货柜签收备查（只登记，不记账）', None),
    (SH_BASE, '资金账户、往来单位、物资档案、参数', None),
    (SH_RULE, '收支类别、二级明细各记到哪个科目', None),
    (SH_RATE, '每月填一次汇率（1 美元＝多少人民币 / ZiG）', None),
    (SH_ACC, '会计科目', None),
    (ECHK, '★ 录完看一眼：有没有 ✗', None),
    (SH_HELP, '怎么录', None),
]
N_HOME_ACC = 6


def build_entry_home(wb, ctx=None):
    if SH_HOME in wb.sheetnames:
        del wb[SH_HOME]
    ws = wb.create_sheet(SH_HOME, 0)
    widths(ws, {'A': 2, 'B': 22, 'C': 12, 'D': 18, 'E': 18, 'F': 4, 'G': 16, 'H': 44})
    title(ws, '=公司名称&" · 录入本"', 'H', C_HOME,
          '💡 平时录账只开这一本，录得快。要看凭证、账簿、报表：先把这本存盘，再打开同一个文件夹里的「A075-2_报表本」，点「更新」'
          '（WPS 是「启用内容」或「更新链接」）。两本别同时开着录。')
    ws.row_dimensions[3].height = 8
    put(ws, 'B4', '资金账户现在的余额', F_SEC, border=False)
    header(ws, 5, [('B', '账户'), ('C', '币种'), ('D', '余额（原币）'), ('E', '折美元（账面）')], C_HOME, height=24)
    K = lambda col, i: f'INDEX({rng(SH_BASE, col, ACC_R0, ACC_R1)},{i})'
    for i in range(1, N_HOME_ACC + 1):
        r = 5 + i
        put(ws, f'B{r}', f'=IF(INDEX(账户名,{i})="","",INDEX(账户名,{i})&"")', F_TXTB, align=AL)
        put(ws, f'C{r}', f'=IF($B{r}="","",INDEX(账户币种,{i})&"")', F_TXT, align=AC)
        put(ws, f'D{r}', f'=IF($B{r}="","",N({K(A_NOWO, i)}))', F_TXT, fmt=MONEY, align=AR_)
        put(ws, f'E{r}', f'=IF($B{r}="","",N({K(A_NOWB, i)}))', F_TXT, fmt=MONEY, align=AR_)
    r = 6 + N_HOME_ACC
    put(ws, f'B{r}', '合计（折美元）', F_TXTB, FILL_TOT, align=AC)
    for c in 'CD':
        put(ws, f'{c}{r}', None, fill_=FILL_TOT)
    put(ws, f'E{r}', f'=SUM(E6:E{r - 1})', F_TXTB, FILL_TOT, MONEY, AR_)
    # 右边：最近一笔、录入检查汇总
    mx = f'MAX({cash_u(K_DATE)})'
    put(ws, 'G5', '最近一笔流水', F_KPI_L, FILL_KPI, align=AC)
    put(ws, 'H5', f'=IF(COUNT({cash_u(K_DATE)})=0,"还没有",YEAR({mx})&"/"&MONTH({mx})&"/"&DAY({mx}))', F_KPI_V, FILL_KPI, align=AC)
    put(ws, 'G6', '录入检查', F_KPI_L, FILL_KPI, align=AC)
    link(put(ws, 'H6', f'={ECHK}!$C$3', F_KPI_V, FILL_KPI, align=AC), ECHK)
    put(ws, 'G7', None, border=False)
    ws.merge_cells('G8:H11')
    put(ws, 'G8', '录的时候只开这一本；看报表打开「A075-2_报表本」。\n别在表头上面插删行、别插删列；漏记的记在最后一行下面，日期写对就行。',
        F_TIP, FILL_TIP, align=ALW)
    # 各表
    r = 8 + N_HOME_ACC
    put(ws, f'B{r}', '各张表（点表名跳过去）', F_SEC, border=False)
    r += 1
    header(ws, r, [('B', '表名'), ('C', '录了几行'), ('D', '✗ 要改'), ('E', '⚠ 看一眼'), ('F', ''), ('G', '做什么用')], C_HOME, height=24)
    ws.merge_cells(f'G{r}:H{r}')
    put(ws, f'H{r}', None, fill_=fill(C_HOME))
    for sh, desc, cols in HOME_TABS:
        if sh not in wb.sheetnames and sh not in (SH_HOME, ECHK):
            continue
        r += 1
        c = put(ws, f'B{r}', sh, Font(name=YH, sz=11, bold=True, color='FF0563C1', underline='single'), align=AL)
        link(c, sh)
        if cols:
            seq, chk = cols
            if sh == SH_CASH:
                n = f'COUNT({cash_u(seq)})'
                x, w = cash_u(chk), cash_u(chk)
            else:
                n = (f'SUMPRODUCT(--((LEN({mod_u(sh, F_NAME)})+LEN({mod_u(sh, F_COST)}))>0))' if sh == SH_FA else
                     f'SUM({mod_u(sh, seq)})' if sh == SH_MAN else f'COUNT({mod_u(sh, seq)})')
                x = w = mod_u(sh, chk)
            put(ws, f'C{r}', f'={n}', F_TXT, fmt='#,##0', align=AC)
            put(ws, f'D{r}', f'=COUNTIF({x},"✗*")', F_RED, fmt='[=0]"-";#,##0', align=AC)
            put(ws, f'E{r}', f'=COUNTIF({w},"⚠*")', F_TXT, fmt='[=0]"-";#,##0', align=AC)
        else:
            for cc in 'CDE':
                put(ws, f'{cc}{r}', None)
        put(ws, f'F{r}', None)
        ws.merge_cells(f'G{r}:H{r}')
        put(ws, f'G{r}', desc, F_TXT, align=AL)
        put(ws, f'H{r}', None)
        ws.row_dimensions[r].height = 20
    ws.freeze_panes = 'A4'
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = C_HOME[2:]
    for s in wb.worksheets:
        s.sheet_view.tabSelected = s is ws
    wb.active = 0
    ensure_names(wb)
    return ws
