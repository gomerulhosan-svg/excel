# -*- coding: utf-8 -*-
"""10/15 新报表（plan.md 第三节，只在报表本；一本版里也建，split 拆本时留在报表本）——照妈妈手工表的样子：
  ① 发生情况报表   ② 记账底稿   ③ 物资出入库台账   ④ 黄金总记录   ⑤ 收支分类报表

入口 build(wb, ctx)：build.make_onebook 在 p_opt 之后、p_wps 之前调，返回日志行。不改别人的表，只读一本版里已有的表 / 列 / 名称：
  现金流水总表 A 日期 B 账户 D 币种 E 类别 G 摘要 L 原币净额 M 折美元 R 二级/备注 X 本账户科目 Z 对方科目 AG 凭证ID AN 类别行；
  _凭证索引「凭证号」；采购入库、领用出库、销售结算、产量登记；基础资料账户、物料档案；记账规则类别；月度汇率；科目表。

写法（README「WPS 存盘修正」一节的规矩）：
  · 成列的公式不写「表名!」、不写往下撑的区域：取录入表第 k 行一律 INDEX(整列名称,k)（common.full_name：流水_M、采购_E…），
    汇总用 cash_u / mod_u（只扫到末行）；别的表的绝对区域（科目表!$A$5:$A$454…）由 p_wps 统一换成名称。
  · 逐笔清单＝「预先排好号」：本表隐藏列里每一笔先算排序键（不符合条件的空着），名次＝COUNTIF(键列,"<="&本笔键)（只有符合条件的几百笔才算），
    清单第 i 行 MATCH(i,名次列,0) 找到是哪一笔（p_opt.step_vx_rank 的路子）。不用 COUNTIFS 套 COUNTIFS。
  · 清单下面的「合计」行跟着笔数走（第 笔数+1 行显示合计），格线、底色用条件格式，空行不画线。
  · 选择格（年 / 月 / 账户 / 物料）空着＝首页报表年月 / 本位币账户 / 第一个物料。报表本里也只有这些选择格。
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from common import *
from s_cash import K_CROW          # 现金流水 AN 列：类别在【记账规则】第几行（找不到＝0，没填＝空）

CUR_B, CUR_CNY, CUR_ZIG = BASE_CUR, 'CNY', 'ZIG'       # 发生情况报表的三栏：美金、人民币、津币（妈妈的三个账户）
NC = CASH_R1 - CASH_R0 + 1                              # 流水 3,960 行
N_RC = RC_R1 - RC_R0 + 1                                # 收支类别 80 个
N_BUY, N_ISS, N_SALE = BUY_R1 - MOD_R0 + 1, ISS_R1 - MOD_R0 + 1, SALE_R1 - MOD_R0 + 1

F_TOTB = Font(name=YH, sz=10, bold=True, color='FF000000')


def _solid(rgb):
    return PatternFill(fill_type='solid', start_color=rgb, end_color=rgb)


CF_TOT = _solid('FFFCE4D6')
CF_SUB = _solid('FFDDEBF7')


def F(sh, col):
    """录入表某列的整列名称（流水_M、采购_E…，p_wps 写进工作簿）"""
    return full_name(sh, col)


def _cf(ws, ref, data_f, total_f=None, head_f=None):
    """清单区：有数的行画格线；合计行加粗、底色；（收支分类）小标题行加粗、浅蓝"""
    if total_f:
        ws.conditional_formatting.add(ref, FormulaRule(formula=[total_f], border=BD, font=F_TOTB, fill=CF_TOT))
    if head_f:
        ws.conditional_formatting.add(ref, FormulaRule(formula=[head_f], border=BD, font=F_TOTB, fill=CF_SUB))
    ws.conditional_formatting.add(ref, FormulaRule(formula=[data_f], border=BD))


def _sel_row(ws, last_col, row=3):
    ws.row_dimensions[row].height = 28
    for i in range(1, CI(last_col) + 1):
        put(ws, f'{CL(i)}{row}', None, fill_=FILL_KPI, border=False)


def _lab(ws, a, text, merge=None):
    if merge:
        ws.merge_cells(f'{a}:{merge}')
    put(ws, a, text, F_KPI_L, FILL_KPI, align=AR_, border=False)


def _inp(ws, a, value=None, fmt=None, merge=None):
    if merge:
        ws.merge_cells(f'{a}:{merge}')
    put(ws, a, value, F_SEL, FILL_SEL, fmt, AC)


def _note(ws, a, value, merge=None):
    if merge:
        ws.merge_cells(f'{a}:{merge}')
    put(ws, a, value, F_NOTE, FILL_KPI, align=AL, border=False)


def _year_dv(ws, a):
    dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2099', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop')
    dv.errorTitle, dv.error = '年份不对', '填四位数的年份，比如 2026；空着就跟【首页】的报表年度'
    dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', '空着＝跟【首页】的报表年度', True
    ws.add_data_validation(dv)
    dv.add(a)


def _month_dv(ws, a):
    dv = DataValidation(type='list', formula1='"1,2,3,4,5,6,7,8,9,10,11,12"', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop')
    dv.errorTitle, dv.error = '月份不对', '从下拉选 1～12；空着就跟【首页】的报表月份'
    dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', '空着＝跟【首页】的报表月份', True
    ws.add_data_validation(dv)
    dv.add(a)


def _list_dv(ws, a, name, tip):
    dv = DataValidation(type='list', formula1=name, allow_blank=True, showErrorMessage=True, errorStyle='stop')
    dv.errorTitle, dv.error = '不在清单里', '请从下拉里选；清单在【基础资料】'
    dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', tip, True
    ws.add_data_validation(dv)
    dv.add(a)


def year_eff(a):
    """选择格 → 实际年份（空着或不像年份＝首页报表年度）"""
    v = f'VALUE({a}&"")'
    return f'=IFERROR(IF(LEN({a}&"")=0,报表年度,IF(AND({v}>=1900,{v}<=9999),INT({v}),报表年度)),报表年度)'


def month_eff(a):
    v = f'VALUE({a}&"")'
    return f'=IFERROR(IF(LEN({a}&"")=0,报表月份,IF(AND({v}>=1,{v}<=12),INT({v}),报表月份)),报表月份)'


def _hdr(ws, a, text, color, merge=None, font=F_HDR):
    if merge:
        ws.merge_cells(f'{a}:{merge}')
        c0, r0 = a.rstrip('0123456789'), int(a.lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))
        c1, r1 = merge.rstrip('0123456789'), int(merge.lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ'))
        for ci in range(CI(c0), CI(c1) + 1):
            for r in range(r0, r1 + 1):
                put(ws, f'{CL(ci)}{r}', None, font, fill(color), align=ACW)
    put(ws, a, text, font, fill(color), align=ACW)


def _print(ws, area, rows=None, landscape=False, one_page=False):
    ws.print_area = area
    if rows:
        ws.print_title_rows = rows
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = 9          # A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, (1 if one_page else 0)
    ws.print_options.horizontalCentered = True


def _style(ws, cols, r0, r1, fmts=None, aligns=None, font=F_TXT):
    """清单区：只设字体 / 格式 / 对齐（格线靠条件格式，空行不画线、印出来不是一页空格子）"""
    fmts, aligns = fmts or {}, aligns or {}
    for r in range(r0, r1 + 1):
        for c in cols:
            x = ws[f'{c}{r}']
            x.font = font
            x.alignment = aligns.get(c, AC)
            if c in fmts:
                x.number_format = fmts[c]


def _helper_head(ws, row, items):
    for c, t in items:
        put(ws, f'{c}{row}', t, F_NOTE, FILL_AUTO, align=ACW)
        ws.column_dimensions[c].width = 11


def _vx_offset(wb, ctx):
    """记账底稿取凭证号：凭证号＝INDEX(凭证号,凭证ID−常数)。流水凭证在 _凭证索引 第 vx0 行起，第 k 个槽位＝凭证ID 100000＋k"""
    import re
    t = wb.defined_names['凭证号'].attr_text
    m = re.match(r"^'?_凭证索引'?!\$" + VX_NO + r"\$(\d+):\$" + VX_NO + r"\$(\d+)$", t)
    assert m, t
    start = int(m.group(1))
    vx0 = ctx['blocks']['流水']['vx0']
    vx = wb[SH_VIDX]
    assert vx[f'{VX_ID}{vx0}'].value == SRC['流水'] * 100000 + 1, vx[f'{VX_ID}{vx0}'].value
    return SRC['流水'] * 100000 - (vx0 - start)


# ═════════════════════════════ ① 发生情况报表 ═════════════════════════════
# 妈妈的月报（u1【发生情况报表】，hand.md §4）：资金来源 / 资金占用 / 资金结余 × 美金、人民币（原币、折美金）、津元（原币、折美金）、美金合计；
# 下面资金余额 期初 / 本期发生 / 期末。只按现金流水算（采购、销售、工资这些不动钱的不进来），每笔按「对方科目」查右边对照表进哪一行：
# 先找整个编码，再找前 6 位、前 4 位；收钱时 / 付钱时可以进不同的行（应收账款：收回＝营业收入，付出＝应收款）。
# 收钱进了「占用」的行（退款、收回借支）记负数，付钱进了「来源」的行（还借款）也记负数。账户之间转钱、换汇只进「资金结余」。
FSQ_SRC = ['股东投资', '股东筹资', '短期借款', '营业收入', '营业外收入', '其他收入']
FSQ_USE = ['应收款', '原材料', '低值易耗品', '固定资产', '在建工程', '无形资产', '生产成本', '长期待摊费用', '应付款', '应交税费',
           '应付职工薪酬', '管理费用', '销售费用', '财务费用', '生活支出', '营业外支出', '其他支出']
FSQ_INT = '（内部转账）'
FSQ_NACC = 5                      # 资金结余列前 5 个账户（她有 3 个），其余的并成「其余账户」一行
FSQ_MAP = [   # (对方科目编码或开头, 收钱时进哪一行, 付钱时进哪一行, 说明)
    ('1001', FSQ_INT, FSQ_INT, '库存现金：账户之间转钱'),
    ('1002', FSQ_INT, FSQ_INT, '银行存款：账户之间转钱'),
    ('1012', FSQ_INT, FSQ_INT, '其他货币资金：账户之间转钱'),
    ('122104', FSQ_INT, FSQ_INT, '资金划转待对冲（内部转账）'),
    ('122105', FSQ_INT, FSQ_INT, '外币兑换待对冲（购汇/结汇）'),
    ('3001', '股东投资', '股东投资', '实收资本'),
    ('3002', '股东投资', '股东投资', '资本公积'),
    ('224101', '股东筹资', '股东筹资', '其他应付款-股东'),
    ('2001', '短期借款', '短期借款', None),
    ('2501', '短期借款', '短期借款', '长期借款'),
    ('224103', '短期借款', '短期借款', '类别「借款」「还款」记在这'),
    ('5001', '营业收入', '营业收入', '主营业务收入（卖黄金）'),
    ('5051', '营业收入', '营业收入', '其他业务收入'),
    ('1122', '营业收入', '应收款', '收回黄金款＝营业收入；付出去＝应收款'),
    ('2203', '营业收入', '营业收入', '预收账款'),
    ('5301', '营业外收入', '营业外收入', None),
    ('5111', '营业外收入', '营业外收入', '投资收益'),
    ('1123', '应收款', '应收款', '预付账款'),
    ('1221', '应收款', '应收款', '其他应收款（借支、保证金）'),
    ('1402', '原材料', '原材料', None),
    ('1403', '原材料', '原材料', None),
    ('1404', '原材料', '原材料', '材料差异（运费、清关费）'),
    ('1405', '原材料', '原材料', None),
    ('540102', '原材料', '原材料', '启用日以前买材料直接算成本'),
    ('1411', '低值易耗品', '低值易耗品', None),
    ('560216', '低值易耗品', '低值易耗品', '管理费用-低值易耗品'),
    ('1601', '固定资产', '固定资产', None),
    ('1604', '在建工程', '在建工程', None),
    ('1701', '无形资产', '无形资产', '采矿权、土地'),
    ('4001', '生产成本', '生产成本', None),
    ('4101', '生产成本', '生产成本', None),
    ('5401', '生产成本', '生产成本', None),
    ('1801', '长期待摊费用', '长期待摊费用', '待摊费用（堆淋、氰化厂）'),
    ('2202', '应付款', '应付款', '应付账款（走采购入库以后付货款）'),
    ('2241', '应付款', '应付款', '其他应付款'),
    ('2232', '应付款', '应付款', '应付利润'),
    ('2211', '应付职工薪酬', '应付职工薪酬', None),
    ('56020901', '应付职工薪酬', '应付职工薪酬', '启用日以前发工资直接记费用'),
    ('2221', '应交税费', '应交税费', None),
    ('5403', '应交税费', '应交税费', '税金及附加（PAYE、土地税…）'),
    ('56021001', '生活支出', '生活支出', '伙食费（食堂、矿泉水、饭费）'),
    ('5602', '管理费用', '管理费用', None),
    ('5601', '销售费用', '销售费用', None),
    ('5603', '财务费用', '财务费用', '手续费、利息'),
    ('5711', '营业外支出', '营业外支出', '罚款等'),
]


def build_fsq(wb, ctx):
    ws = wb.create_sheet(SH_FSQ)
    widths(ws, {'A': 6, 'B': 5, 'C': 24, 'D': 16, 'E': 16, 'F': 15, 'G': 16, 'H': 15, 'I': 17, 'J': 2,
                'K': 10, 'L': 13, 'M': 13, 'N': 30})
    Y, M_, M0, M1 = '$R$3', '$S$3', '$T$3', '$U$3'
    title(ws, f'={Y}&"年"&{M_}&"月发生情况报表"', 'I', C_RPT,
          '💡 跟你手工的「发生情况报表」一样：这个月的钱从哪来（资金来源）、花到哪去（资金占用）、各账户增减（资金结余），'
          '下面是三种钱的月初、本月、月末余额。只按【现金流水总表】算，每笔按「对方科目」进哪一行见右边对照表；'
          '账户之间转钱、换汇只进资金结余。年份、月份空着就跟【首页】。', h2=48)
    _sel_row(ws, 'I')
    _lab(ws, 'B3', '年份：', 'C3')
    _inp(ws, 'D3')
    _lab(ws, 'E3', '月份：')
    _inp(ws, 'F3')
    _note(ws, 'G3', f'="→ "&{Y}&"年"&{M_}&"月"&IF(LEN($D$3&"")+LEN($F$3&"")=0,"（空着＝跟首页）","")', 'I3')
    _year_dv(ws, 'D3')
    _month_dv(ws, 'F3')
    ws['R3'], ws['S3'] = year_eff('$D$3'), month_eff('$F$3')
    ws['T3'], ws['U3'] = f'=DATE({Y},{M_},1)', f'=DATE({Y},{M_}+1,1)'
    for c in 'RSTU':
        ws[f'{c}3'].font = F_NOTE
    ws['T3'].number_format = ws['U3'].number_format = DATE

    # 表头（两行，跟手工表一样）
    _hdr(ws, 'A4', '类别', C_RPT, 'A5')
    _hdr(ws, 'B4', '序号', C_RPT, 'B5')
    _hdr(ws, 'C4', '科目', C_RPT, 'C5')
    _hdr(ws, 'D4', '美金', C_RPT, 'D5')
    _hdr(ws, 'E4', '人民币', C_RPT, 'F4')
    _hdr(ws, 'E5', '原币', C_RPT)
    _hdr(ws, 'F5', '折美金', C_RPT)
    _hdr(ws, 'G4', '津币 ZiG', C_RPT, 'H4')
    _hdr(ws, 'G5', '原币', C_RPT)
    _hdr(ws, 'H5', '折美金', C_RPT)
    _hdr(ws, 'I4', '美金合计', C_RPT, 'I5')
    ws.row_dimensions[4].height = ws.row_dimensions[5].height = 20

    # 对照表（右边，灰格，生成的；要改找做表的人改 s_extra.FSQ_MAP）
    MR0 = 6
    MR1 = MR0 + len(FSQ_MAP) - 1
    ws.merge_cells('K3:N3')
    put(ws, 'K3', '对照表：流水的对方科目 → 进哪一行（先找整个编码，再找前 6 位、前 4 位；都没有＝其他收入 / 其他支出）',
        F_NOTE, FILL_AUTO, align=ALW)
    ws.row_dimensions[3].height = 30
    for c, t in zip('KLMN', ('对方科目\n编码 / 开头', '收钱时\n进哪一行', '付钱时\n进哪一行', '说明')):
        _hdr(ws, f'{c}4', t, C_BASE, f'{c}5', font=F_HDR)
    rows_ok = set(FSQ_SRC) | set(FSQ_USE) | {FSQ_INT}
    for i, (code, a, b, note) in enumerate(FSQ_MAP):
        assert a in rows_ok and b in rows_ok, (code, a, b)
        r = MR0 + i
        for c, v in zip('KLMN', (code, a, b, note)):
            put(ws, f'{c}{r}', v, F_NOTE, FILL_AUTO, '@' if c == 'K' else None, AL if c == 'N' else AC)
    mk, mt = f'$K${MR0}:$K${MR1}', f'$L${MR0}:$M${MR1}'

    # 隐藏 P 列：流水第 k 笔（P 列第 k+5 行，跟流水同一行）进哪一行
    L = f'INDEX({F(SH_CASH, K_NET)},ROW()-{CASH_R0 - 1})'
    Z = f'INDEX({F(SH_CASH, K_CCODE)},ROW()-{CASH_R0 - 1})&""'
    pos = f'IFERROR(MATCH({Z},{mk},0),IFERROR(MATCH(LEFT({Z},6),{mk},0),MATCH(LEFT({Z},4),{mk},0)))'
    oth = f'IF({L}>0,"{FSQ_SRC[-1]}","{FSQ_USE[-1]}")'
    _helper_head(ws, CASH_R0 - 1, [('P', '流水这一笔\n进哪一行')])
    for r in range(CASH_R0, CASH_R1 + 1):
        ws[f'P{r}'] = f'=IF(N({L})=0,"",IF({Z}="",{oth},IFERROR(INDEX({mt},{pos},IF({L}>0,1,2)),{oth})))'
    HU = f'$P${CASH_R0}:INDEX($P${CASH_R0}:$P${CASH_R1},流水末行)'      # 跟 cash_u 一样只到流水末行，跟流水各列一样大
    dates = f'{cash_u(K_DATE)},">="&{M0},{cash_u(K_DATE)},"<"&{M1}'

    def sumifs(col, cur=None, lab=None, extra=''):
        p = [cash_u(col)]
        if lab is not None:
            p += [HU, lab]
        if extra:
            p += [extra]
        if cur:
            p += [cash_u(K_CUR), f'"{cur}"']
        return 'SUMIFS(' + ','.join(p) + ',' + dates + ')'

    VAL = [('D', K_UZS, CUR_B), ('E', K_NET, CUR_CNY), ('F', K_UZS, CUR_CNY), ('G', K_NET, CUR_ZIG), ('H', K_UZS, CUR_ZIG),
           ('I', K_UZS, None)]

    def block(r0, label, items, sign):
        r = r0
        for i, it in enumerate(items):
            put(ws, f'B{r}', i + 1, F_TXT, align=AC)
            put(ws, f'C{r}', it, F_TXT, align=AL)
            for c, col, cur in VAL:
                put(ws, f'{c}{r}', f'={sign}{sumifs(col, cur, f"$C{r}")}', F_TXT, fmt=MONEY, align=AR_)
            r += 1
        rt = r
        ws.merge_cells(f'B{rt}:C{rt}')
        put(ws, f'B{rt}', '合　　计', F_TXTB, FILL_TOT, align=AC)
        put(ws, f'C{rt}', None, F_TXTB, FILL_TOT)
        for c, _, _ in VAL:
            put(ws, f'{c}{rt}', f'=SUM({c}{r0}:{c}{rt - 1})', F_TXTB, FILL_TOT, MONEY, AR_)
        ws.merge_cells(f'A{r0}:A{rt}')
        put(ws, f'A{r0}', label, F_TXTB, FILL_SUBH, align=ACW)
        for rr in range(r0 + 1, rt + 1):
            put(ws, f'A{rr}', None, F_TXTB, FILL_SUBH)
        return rt

    r_src0 = 6
    r_src_t = block(r_src0, '资\n金\n来\n源', FSQ_SRC, '')
    r_use0 = r_src_t + 1
    r_use_t = block(r_use0, '资\n金\n占\n用', FSQ_USE, '-')

    # 资金结余：各账户本月增减（含账户之间转钱）
    r_acc0 = r_use_t + 1
    for k in range(1, FSQ_NACC + 1):
        r = r_acc0 + k - 1
        nm = f'INDEX(账户名,{k})'
        put(ws, f'B{r}', k, F_TXT, align=AC)
        put(ws, f'C{r}', f'=IF({nm}&""="","",IFERROR(INDEX({acc(AC_NAME)},MATCH(LEFT(INDEX(账户科目,{k})&"",4),{acc(AC_CODE)},0))&"-","")'
                         f'&{nm})', F_TXT, align=AL)
        for c, col, cur in VAL:
            put(ws, f'{c}{r}', f'=IF({nm}&""="","",{sumifs(col, cur, extra=cash_u(K_ACC) + "," + nm)})', F_TXT, fmt=MONEY, align=AR_)
    r_oth = r_acc0 + FSQ_NACC
    put(ws, f'B{r_oth}', FSQ_NACC + 1, F_TXT, align=AC)
    put(ws, f'C{r_oth}', '其余账户', F_TXT, align=AL)
    for c, col, cur in VAL:
        put(ws, f'{c}{r_oth}', f'={sumifs(col, cur)}-SUM({c}{r_acc0}:{c}{r_oth - 1})', F_TXT, fmt=MONEY, align=AR_)
    r_acc_t = r_oth + 1
    ws.merge_cells(f'B{r_acc_t}:C{r_acc_t}')
    put(ws, f'B{r_acc_t}', '合　　计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'C{r_acc_t}', None, F_TXTB, FILL_TOT)
    for c, _, _ in VAL:
        put(ws, f'{c}{r_acc_t}', f'=SUM({c}{r_acc0}:{c}{r_oth})', F_TXTB, FILL_TOT, MONEY, AR_)
    ws.merge_cells(f'A{r_acc0}:A{r_acc_t}')
    put(ws, f'A{r_acc0}', '资\n金\n结\n余', F_TXTB, FILL_SUBH, align=ACW)
    for rr in range(r_acc0 + 1, r_acc_t + 1):
        put(ws, f'A{rr}', None, F_TXTB, FILL_SUBH)
    r_int = r_acc_t + 1
    ws.merge_cells(f'A{r_int}:C{r_int}')
    put(ws, f'A{r_int}', '其中：账户之间转钱、换汇（不算来源、占用）', F_NOTE, align=ALW)
    for c in 'BC':
        put(ws, f'{c}{r_int}', None, F_NOTE)
    for c, col, cur in VAL:
        put(ws, f'{c}{r_int}', f'={sumifs(col, cur, chr(34) + FSQ_INT + chr(34))}', F_NOTE, fmt=MONEY, align=AR_)
    ws.row_dimensions[r_int].height = 30

    # 资金余额：期初 / 本期发生 / 期末（折美金用账面美元，跟流水 O 列一样按当时汇率累计，不按月末重估）
    rh = r_int + 2
    _hdr(ws, f'A{rh}', '项目', C_RPT, f'B{rh + 1}')
    _hdr(ws, f'C{rh}', '科目', C_RPT, f'C{rh + 1}')
    _hdr(ws, f'D{rh}', '美金', C_RPT, f'D{rh + 1}')
    _hdr(ws, f'E{rh}', '人民币', C_RPT, f'F{rh}')
    _hdr(ws, f'E{rh + 1}', '原币', C_RPT)
    _hdr(ws, f'F{rh + 1}', '折美金', C_RPT)
    _hdr(ws, f'G{rh}', '津币 ZiG', C_RPT, f'H{rh}')
    _hdr(ws, f'G{rh + 1}', '原币', C_RPT)
    _hdr(ws, f'H{rh + 1}', '折美金', C_RPT)
    _hdr(ws, f'I{rh}', '各项美金\n合计', C_RPT, f'I{rh + 1}')
    ws.row_dimensions[rh].height = ws.row_dimensions[rh + 1].height = 20
    r_o, r_m, r_c = rh + 2, rh + 3, rh + 4
    before = f'{cash_u(K_DATE)},"<"&{M0}'

    def opening(col, cur, base):
        if cur is None:
            return f'=SUM(账户期初本位)+SUMIFS({cash_u(col)},{before})'
        return f'=SUMIF(账户币种,"{cur}",{base})+SUMIFS({cash_u(col)},{cash_u(K_CUR)},"{cur}",{before})'

    for r, t in ((r_o, '期初余额'), (r_m, '本期发生'), (r_c, '期末结余')):
        put(ws, f'C{r}', t, F_TXTB, align=AC)
        for c, col, cur in VAL:
            if r == r_o:
                v = opening(col, cur, '账户期初' if col == K_NET or cur == CUR_B else '账户期初本位')
            elif r == r_m:
                v = f'={sumifs(K_NET if (c == "D" or col == K_NET) else col, cur)}'
            else:
                v = f'={c}{r_o}+{c}{r_m}'
            put(ws, f'{c}{r}', v, F_TXTB if r == r_c else F_TXT, FILL_TOT if r == r_c else None, MONEY, AR_)
    ws.merge_cells(f'A{r_o}:B{r_c}')
    put(ws, f'A{r_o}', '资金\n余额', F_TXTB, FILL_SUBH, align=ACW)
    for rr in range(r_o, r_c + 1):
        for c in 'AB':
            if (c, rr) != ('A', r_o):
                put(ws, f'{c}{rr}', None, F_TXTB, FILL_SUBH)
    r_chk = r_c + 2
    d1 = f'I{r_src_t}-I{r_use_t}+I{r_int}-I{r_acc_t}'
    d2 = f'I{r_acc_t}-I{r_m}'
    ws.merge_cells(f'A{r_chk}:I{r_chk}')
    put(ws, f'A{r_chk}', f'=IF(AND(ROUND({d1},2)=0,ROUND({d2},2)=0),"√ 对上了：来源 − 占用 ＋ 账户之间转钱 ＝ 资金结余 ＝ 本期发生",'
                         f'"✗ 对不上：差 "&FIXED({d1},2)&" / "&FIXED({d2},2)&"（流水有没填日期、账户的，或对照表有对不上的行）")',
        F_RED, align=AL, border=False)
    for r in range(6, r_c + 1):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = 20
    hide(ws, list('OPQRSTU'))
    ws.freeze_panes = 'D6'
    _print(ws, f'A1:I{r_chk}', landscape=False, one_page=True)
    ctx['fsq_rows'] = dict(src=(r_src0, r_src_t), use=(r_use0, r_use_t), acc=(r_acc0, r_acc_t), int=r_int,
                           bal=(r_o, r_m, r_c), chk=r_chk)
    return f'【{SH_FSQ}】来源 {len(FSQ_SRC)} 行、占用 {len(FSQ_USE)} 行、结余 {FSQ_NACC + 1} 个账户行、资金余额 3 行；对照表 {len(FSQ_MAP)} 条'


# ═════════════════════════════ ② 记账底稿 ═════════════════════════════
# 妈妈的「美元 / 人民币 / FBC 记帐底稿帐」（本月凭证汇总记录）：选账户、月份，这个账户本月每一笔一行，
# 月、日、凭证号（＝【凭证汇总】的当月第几号）、摘要、借方科目（一级、二级）、贷方科目（一级、二级）、折美金、原币，按凭证号排，最后合计。
DG_N = 600                        # 清单容量（她 2022 年美元最多的一个月 537 行）
DG_R0 = 7                         # 第 7 行起列清单；隐藏 L、M 列第 7 行起每行对着流水第 k 笔（k＝行号−6）


def build_dg(wb, ctx):
    ws = wb.create_sheet(SH_DG)
    widths(ws, {'A': 5, 'B': 5, 'C': 7, 'D': 40, 'E': 13, 'F': 20, 'G': 13, 'H': 20, 'I': 15, 'J': 15, 'K': 2})
    ACC, Y, M_, M0, M1, CNT = '$L$3', '$M$3', '$N$3', '$O$3', '$P$3', '$Q$3'
    TOTM, TOTL, INC, EXP, OPEN, NET = '$R$3', '$S$3', '$T$3', '$U$3', '$V$3', '$W$3'
    R1 = DG_R0 + DG_N - 1
    K1 = DG_R0 + NC - 1               # 流水对应的隐藏行到第 3966 行
    off = _vx_offset(wb, ctx)
    title(ws, f'={ACC}&"记账底稿（"&{Y}&"年"&{M_}&"月 · 本月凭证汇总记录）"', 'J', C_LED,
          '💡 跟你手工的「记帐底稿帐」一样：选账户、年份、月份，这个账户这个月每一笔一行，按凭证号排（凭证号＝【凭证汇总】里的当月第几号；'
          '同一天同一个账户的收款、付款各并成一张凭证）。借方、贷方科目是自动记账记到的科目；摘要后面括号里是你写的「二级/备注」。'
          '年份、月份空着就跟【首页】。打印只印有数的那几页就行（后面是空行）。', h2=48)
    _sel_row(ws, 'J')
    _lab(ws, 'A3', '账户：', 'B3')
    names = [wb[SH_BASE][f'{A_NAME}{r}'].value for r in range(ACC_R0, ACC_R1 + 1)]
    _inp(ws, 'C3', '美元' if '美元' in names else None, merge='D3')
    _list_dv(ws, 'C3', '账户列表', '选账户（美元 / 人民币 / FBC）；空着＝美元账户')
    _lab(ws, 'E3', '年份：')
    _inp(ws, 'F3')
    _lab(ws, 'G3', '月份：')
    _inp(ws, 'H3')
    _year_dv(ws, 'F3')
    _month_dv(ws, 'H3')
    _note(ws, 'I3', f'=IF(LEN($F$3&"")+LEN($H$3&"")=0,"年月空着＝跟首页","")', 'J3')
    ws['L3'] = f'=IF(TRIM($C$3&"")="",IFERROR(INDEX(账户名,MATCH(本位币,账户币种,0))&"",INDEX(账户名,1)&""),TRIM($C$3&""))'
    ws['M3'], ws['N3'] = year_eff('$F$3'), month_eff('$H$3')
    ws['O3'], ws['P3'] = f'=DATE({Y},{M_},1)', f'=DATE({Y},{M_}+1,1)'
    ws['Q3'] = f'=COUNT($M${DG_R0}:INDEX($M${DG_R0}:$M${K1},流水末行))'
    acc_ = f'{cash_u(K_ACC)},{ACC}'
    dates = f'{cash_u(K_DATE)},">="&{M0},{cash_u(K_DATE)},"<"&{M1}'
    ws['T3'] = f'=SUMIFS({cash_u(K_UZS)},{acc_},{cash_u(K_UZS)},">0",{dates})'
    ws['U3'] = f'=-SUMIFS({cash_u(K_UZS)},{acc_},{cash_u(K_UZS)},"<0",{dates})'
    ws['R3'] = '=$T$3+$U$3'
    ws['S3'] = (f'=SUMIFS({cash_u(K_NET)},{acc_},{cash_u(K_NET)},">0",{dates})'
                f'-SUMIFS({cash_u(K_NET)},{acc_},{cash_u(K_NET)},"<0",{dates})')
    ws['V3'] = f'=SUMIF(账户名,{ACC},账户期初)+SUMIFS({cash_u(K_NET)},{acc_},{cash_u(K_DATE)},"<"&{M0})'
    ws['W3'] = f'=SUMIFS({cash_u(K_NET)},{acc_},{dates})'
    for c in 'LMNOPQRSTUVW':
        ws[f'{c}3'].font = F_NOTE
    ws['O3'].number_format = ws['P3'].number_format = DATE
    ws.merge_cells('A4:J4')
    put(ws, 'A4', f'="本月 "&{CNT}&" 笔　收入 "&FIXED({INC},2)&"　支出 "&FIXED({EXP},2)&"（折美金）　｜　"&{ACC}&"　月初余额 "'
                  f'&FIXED({OPEN},2)&"　月末余额 "&FIXED({OPEN}+{NET},2)&"（原币）"'
                  f'&IF({CNT}>{DG_N},"　✗ 超过 {DG_N} 笔，只列了前 {DG_N} 笔","")', F_TXTB, FILL_SUBH, align=AL, border=False)
    ws.row_dimensions[4].height = 22
    # 表头两行（跟手工表一样）
    _hdr(ws, 'A5', f'={Y}&"年"', C_LED, 'B5')
    _hdr(ws, 'A6', '月', C_LED)
    _hdr(ws, 'B6', '日', C_LED)
    _hdr(ws, 'C5', '凭证号', C_LED, 'C6')
    _hdr(ws, 'D5', '摘要', C_LED, 'D6')
    _hdr(ws, 'E5', '借方科目', C_LED, 'F5')
    _hdr(ws, 'E6', '一级科目', C_LED)
    _hdr(ws, 'F6', '二级科目', C_LED)
    _hdr(ws, 'G5', '贷方科目', C_LED, 'H5')
    _hdr(ws, 'G6', '一级科目', C_LED)
    _hdr(ws, 'H6', '二级科目', C_LED)
    _hdr(ws, 'I5', '折美金金额', C_LED)
    _hdr(ws, 'I6', '按当时汇率', C_LED)
    _hdr(ws, 'J5', '原币金额', C_LED, 'J6')
    ws.row_dimensions[5].height = ws.row_dimensions[6].height = 20
    _helper_head(ws, DG_R0 - 1, [('L', '排序键\n(流水第k笔)'), ('M', '名次'), ('N', '第i行是\n流水第几笔'), ('O', '借方编码'), ('P', '贷方编码')])
    # 隐藏 L / M：流水第 k 笔（第 k+6 行）——是这个账户、这个月的才有排序键（日×10^9＋凭证ID×10^4＋k：按凭证号、同一张里按录入顺序）
    k_ = f'ROW()-{DG_R0 - 1}'
    X = lambda col: f'INDEX({F(SH_CASH, col)},{k_})'
    for r in range(DG_R0, K1 + 1):
        ws[f'L{r}'] = (f'=IF({X(K_ACC)}&""<>{ACC},"",IF(N({X(K_NET)})=0,"",IF(NOT(ISNUMBER({X(K_DATE)})),"",'
                       f'IF(OR({X(K_DATE)}<{M0},{X(K_DATE)}>={M1}),"",DAY({X(K_DATE)})*10^9'
                       f'+IF(ISNUMBER({X(K_VID)}),{X(K_VID)},199999)*10^4+{k_}))))')
        ws[f'M{r}'] = f'=IF($L{r}="","",COUNTIF($L${DG_R0}:INDEX($L${DG_R0}:$L${K1},流水末行),"<="&$L{r}))'
    # 清单第 i 行（第 i+6 行）
    i_ = f'ROW()-{DG_R0 - 1}'
    T = f'{i_}={CNT}+1'
    P = lambda col, r: f'INDEX({F(SH_CASH, col)},$N{r})'
    for r in range(DG_R0, R1 + 1):
        ws[f'N{r}'] = f'=IF({i_}>{CNT},"",MATCH({i_},$M${DG_R0}:INDEX($M${DG_R0}:$M${K1},流水末行),0))'
        ws[f'O{r}'] = f'=IF($N{r}="","",IF({P(K_NET, r)}>0,{P(K_ACODE, r)},{P(K_CCODE, r)})&"")'
        ws[f'P{r}'] = f'=IF($N{r}="","",IF({P(K_NET, r)}>0,{P(K_CCODE, r)},{P(K_ACODE, r)})&"")'
        ws[f'A{r}'] = f'=IF($N{r}="",IF({T},"合计",""),MONTH({P(K_DATE, r)}))'
        ws[f'B{r}'] = f'=IF($N{r}="","",DAY({P(K_DATE, r)}))'
        ws[f'C{r}'] = f'=IF($N{r}="","",IF(ISNUMBER({P(K_VID, r)}),INDEX(凭证号,{P(K_VID, r)}-{off}),""))'
        ws[f'D{r}'] = (f'=IF($N{r}="","",{P(K_MEMO, r)}&IF({P(K_NOTE, r)}&""="","","（"&{P(K_NOTE, r)}&"）"))')
        for cn, c2, cc in (('E', 'F', 'O'), ('G', 'H', 'P')):
            ws[f'{cn}{r}'] = (f'=IF($N{r}="","",IF(${cc}{r}="","（没记账）",'
                              f'IFERROR(INDEX({acc(AC_NAME)},MATCH(LEFT(${cc}{r},4),{acc(AC_CODE)},0)),${cc}{r})))')
            ws[f'{c2}{r}'] = (f'=IF(LEN(${cc}{r})<=4,"",IFERROR(MID(INDEX({acc(AC_FULL)},MATCH(${cc}{r},{acc(AC_CODE)},0)),'
                              f'LEN(${cn}{r})+2,99),""))')
        ws[f'I{r}'] = f'=IF($N{r}="",IF({T},{TOTM},""),ABS(N({P(K_UZS, r)})))'
        ws[f'J{r}'] = f'=IF($N{r}="",IF({T},{TOTL},""),ABS(N({P(K_NET, r)})))'
    _style(ws, 'ABCDEFGHIJ', DG_R0, R1, fmts={'C': '0', 'I': MONEY0, 'J': MONEY0},
           aligns={'D': AL, 'E': AL, 'F': AL, 'G': AL, 'H': AL, 'I': AR_, 'J': AR_})
    for r in range(DG_R0, K1 + 1):
        for c in 'LM':
            ws[f'{c}{r}'].font = F_NOTE
    _cf(ws, f'A{DG_R0}:J{R1}', f'$N{DG_R0}<>""', f'AND($N{DG_R0}="",ROW()-{DG_R0 - 1}={CNT}+1)')
    hide(ws, list('LMNOPQRSTUVW'))
    ws.freeze_panes = f'D{DG_R0}'
    _print(ws, f'A1:J{R1}', '5:6', landscape=True)
    return f'【{SH_DG}】清单 {DG_N} 行（流水 {NC:,} 笔各排一个号，选中账户、月份的才算名次）'


# ═════════════════════════════ ③ 物资出入库台账 ═════════════════════════════
# 妈妈的「XX 出入库台账」（u2）：选物资、年份，入库（【采购入库】）、出库（【领用出库】）按日期逐笔排（同一天先入后出），结余自动滚；
# 期初＝物料档案期初数量＋这一年以前的入库−出库。
WZ_N = 500
WZ_R0 = 8                         # 第 7 行期初，第 8 行起清单；隐藏 I、J 列第 8 行起：采购第 1～600 行、领用第 1～1000 行


def build_wz(wb, ctx):
    ws = wb.create_sheet(SH_WZ)
    widths(ws, {'A': 5, 'B': 5, 'C': 44, 'D': 13, 'E': 13, 'F': 13, 'G': 18, 'H': 2})
    MAT, Y, CNT, TIN, TOUT, OPEN, UNIT = '$I$3', '$J$3', '$K$3', '$L$3', '$M$3', '$N$3', '$O$3'
    R1 = WZ_R0 + WZ_N - 1
    NB, NI = N_BUY, N_ISS
    H1 = WZ_R0 + NB + NI - 1
    first = wb[SH_BASE][f'{M_NAME}{MAT_R0}'].value
    mname = rng(SH_BASE, M_NAME, MAT_R0, MAT_R1)
    title(ws, f'={MAT}&"出入库台账（"&{Y}&"年）"', 'G', C_OPS,
          '💡 跟你手工的「出入库台账」一样：选物资、年份，入库取【采购入库】、出库取【领用出库】，按日期一笔一行（同一天先入后出），'
          '结余自动算，不用手算。期初＝【基础资料·物料档案】的期初数量加上这一年以前的入库、减出库。年份空着就跟【首页】。', h2=40)
    _sel_row(ws, 'G')
    _lab(ws, 'A3', '物资：', 'B3')
    _inp(ws, 'C3', first)
    _list_dv(ws, 'C3', '物料列表', f'选物资（清单在【基础资料·物料档案】）；空着＝第一个（{first}）')
    _lab(ws, 'D3', '年份：')
    _inp(ws, 'E3')
    _year_dv(ws, 'E3')
    _note(ws, 'F3', f'="单位："&{UNIT}&IF(LEN($E$3&"")=0,"　年份空着＝跟首页","")', 'G3')
    ws['I3'] = f'=IF(TRIM($C$3&"")="",INDEX({mname},1)&"",TRIM($C$3&""))'
    ws['J3'] = year_eff('$E$3')
    ws['K3'] = f'=COUNT($J${WZ_R0}:$J${H1})'
    y0, y1 = f'DATE({Y},1,1)', f'DATE({Y}+1,1,1)'
    bq, bm, bd = mod_u(SH_BUY, B_QTY), mod_u(SH_BUY, B_MAT), mod_u(SH_BUY, B_DATE)
    iq, im, idt = mod_u(SH_ISS, I_QTY), mod_u(SH_ISS, I_MAT), mod_u(SH_ISS, I_DATE)
    ws['L3'] = f'=SUMIFS({bq},{bm},{MAT},{bd},">="&{y0},{bd},"<"&{y1})'
    ws['M3'] = f'=SUMIFS({iq},{im},{MAT},{idt},">="&{y0},{idt},"<"&{y1})'
    ws['N3'] = (f'=IFERROR(N(INDEX({rng(SH_BASE, M_Q0, MAT_R0, MAT_R1)},MATCH({MAT},{mname},0))),0)'
                f'+SUMIFS({bq},{bm},{MAT},{bd},"<"&{y0})-SUMIFS({iq},{im},{MAT},{idt},"<"&{y0})')
    ws['O3'] = f'=IFERROR(INDEX({rng(SH_BASE, M_UNIT, MAT_R0, MAT_R1)},MATCH({MAT},{mname},0))&"","")'
    for c in 'IJKLMNO':
        ws[f'{c}3'].font = F_NOTE
    spec = f'IFERROR(INDEX({rng(SH_BASE, M_SPEC, MAT_R0, MAT_R1)},MATCH({MAT},{mname},0))&"","")'
    cat = f'IFERROR(INDEX({rng(SH_BASE, M_CAT, MAT_R0, MAT_R1)},MATCH({MAT},{mname},0))&"","")'
    ws.merge_cells('A4:G4')
    put(ws, 'A4', f'=IF(COUNTIF({mname},{MAT})=0,"✗ 物料档案里没有「"&{MAT}&"」",'
                  f'"规格："&IF({spec}="","-",{spec})&"　类别："&{cat}&"　本年入库 "&{TIN}&"　出库 "&{TOUT}&"　年末结余 "&({OPEN}+{TIN}-{TOUT})'
                  f'&IF({CNT}>{WZ_N},"　✗ 超过 {WZ_N} 笔，只列了前 {WZ_N} 笔",""))', F_TXTB, FILL_SUBH, align=AL, border=False)
    ws.row_dimensions[4].height = 22
    _hdr(ws, 'A5', f'={Y}&"年"', C_OPS, 'B5')
    _hdr(ws, 'A6', '月', C_OPS)
    _hdr(ws, 'B6', '日', C_OPS)
    _hdr(ws, 'C5', '摘要（进货单位 · 用途 · 堆号 · 单号）', C_OPS, 'C6')
    _hdr(ws, 'D5', f'="入库（"&{UNIT}&"）"', C_OPS, 'D6')
    _hdr(ws, 'E5', f'="出库（"&{UNIT}&"）"', C_OPS, 'E6')
    _hdr(ws, 'F5', '结余', C_OPS, 'F6')
    _hdr(ws, 'G5', '领取人/备注', C_OPS, 'G6')
    ws.row_dimensions[5].height = ws.row_dimensions[6].height = 20
    r0 = WZ_R0 - 1
    put(ws, f'C{r0}', f'=IF({Y}<=YEAR(起始月份),"期初结余（建账）","上年结转")', F_TXTB, FILL_SUBH, align=AL)
    put(ws, f'F{r0}', f'={OPEN}', F_TXTB, FILL_SUBH, 'General', AR_)
    for c in 'ABDEG':
        put(ws, f'{c}{r0}', None, F_TXTB, FILL_SUBH)
    _helper_head(ws, WZ_R0 - 1, [('I', '排序键\n(采购/领用)'), ('J', '名次'), ('K', '第i行是\n第几个'), ('L', '日期')])
    # 隐藏 I / J：第 p 个（p≤600 采购第 p 行，否则领用第 p−600 行）——是这个物资、这一年、日期和数量都是数的才有排序键（天×10^4＋p）
    for r in range(WZ_R0, H1 + 1):
        p = r - WZ_R0 + 1
        if p <= NB:
            k_, sh, cm, cd, cq = f'ROW()-{WZ_R0 - 1}', SH_BUY, B_MAT, B_DATE, B_QTY
        else:
            k_, sh, cm, cd, cq = f'ROW()-{WZ_R0 - 1 + NB}', SH_ISS, I_MAT, I_DATE, I_QTY
        X = lambda col: f'INDEX({F(sh, col)},{k_})'
        ws[f'I{r}'] = (f'=IF({X(cm)}&""<>{MAT},"",IF(NOT(ISNUMBER({X(cd)})),"",IF(NOT(ISNUMBER({X(cq)})),"",'
                       f'IF(YEAR({X(cd)})<>{Y},"",(INT({X(cd)})-{y0})*10^4+ROW()-{WZ_R0 - 1}))))')
        ws[f'J{r}'] = f'=IF($I{r}="","",COUNTIF($I${WZ_R0}:$I${H1},"<="&$I{r}))'
    i_ = f'ROW()-{WZ_R0 - 1}'
    T = f'{i_}={CNT}+1'
    B = lambda col, r: f'INDEX({F(SH_BUY, col)},$K{r})'
    I = lambda col, r: f'INDEX({F(SH_ISS, col)},$K{r}-{NB})'
    nz = lambda x: f'IF({x}&""="","","　"&{x})'
    for r in range(WZ_R0, R1 + 1):
        ws[f'K{r}'] = f'=IF({i_}>{CNT},"",MATCH({i_},$J${WZ_R0}:$J${H1},0))'
        ws[f'L{r}'] = f'=IF($K{r}="","",IF($K{r}<={NB},{B(B_DATE, r)},{I(I_DATE, r)}))'
        ws[f'A{r}'] = f'=IF($K{r}="",IF({T},"合计",""),MONTH($L{r}))'
        ws[f'B{r}'] = f'=IF($K{r}="","",DAY($L{r}))'
        ws[f'C{r}'] = (f'=IF($K{r}="","",IF($K{r}<={NB},"进货"&{nz(B(B_SUP, r))}&IF({B(B_NO, r)}&""="","","　单号 "&{B(B_NO, r)})'
                       f'&{nz(B(B_NOTE, r))},{I(I_USE, r)}&{nz(I(I_PRJ, r))}&IF({I(I_NO, r)}&""="","","　单号 "&{I(I_NO, r)})))')
        ws[f'D{r}'] = f'=IF($K{r}="",IF({T},{TIN},""),IF($K{r}<={NB},{B(B_QTY, r)},""))'
        ws[f'E{r}'] = f'=IF($K{r}="",IF({T},{TOUT},""),IF($K{r}<={NB},"",{I(I_QTY, r)}))'
        ws[f'F{r}'] = f'=IF($K{r}="",IF({T},F{r - 1},""),F{r - 1}+N(D{r})-N(E{r}))'
        ws[f'G{r}'] = f'=IF($K{r}="","",IF($K{r}<={NB},"",{I(I_NOTE, r)}&""))'
    _style(ws, 'ABCDEFG', WZ_R0, R1, fmts={'D': QTY, 'E': QTY, 'F': 'General'},
           aligns={'C': AL, 'G': AL, 'D': AR_, 'E': AR_, 'F': AR_})
    for r in range(WZ_R0, H1 + 1):
        for c in 'IJ':
            ws[f'{c}{r}'].font = F_NOTE
    _cf(ws, f'A{WZ_R0}:G{R1}', f'$K{WZ_R0}<>""', f'AND($K{WZ_R0}="",ROW()-{WZ_R0 - 1}={CNT}+1)')
    hide(ws, list('IJKLMNO'))
    ws.freeze_panes = f'A{WZ_R0}'
    _print(ws, f'A1:G{R1}', '5:6')
    return f'【{SH_WZ}】清单 {WZ_N} 行（采购 {NB} 行＋领用 {NI} 行各排一个号）；默认物资「{first}」'


# ═════════════════════════════ ④ 黄金总记录 ═════════════════════════════
# 妈妈的「黄金总记录」（u1【部分收入】）：上面按月 冶炼克数（【产量登记】）、出售克数（【销售结算】）、差额；
# 下面逐笔出售：出售日期、买方、克数、单价、金额、累计总收入、当月人民币汇率（【月度汇率】CNY）、折人民币。
GD_R0 = 22                        # 逐笔清单第 22 行起（300 行，跟销售结算一样多）；隐藏 L、M 列同一行对着销售第 k 行


def build_gold(wb, ctx):
    ws = wb.create_sheet(SH_GOLD)
    widths(ws, {'A': 8, 'B': 14, 'C': 15, 'D': 16, 'E': 14, 'F': 17, 'G': 17, 'H': 12, 'I': 17, 'J': 18, 'K': 2})
    Y, CNT, TG, TU, TR = '$L$3', '$M$3', '$N$3', '$O$3', '$P$3'
    R1 = GD_R0 + N_SALE - 1
    title(ws, f'="黄金总记录（"&{Y}&"年）"', 'J', C_SALE,
          '💡 跟你手工的「黄金总记录」一样：上面按月看冶炼了多少克（【产量登记】）、卖了多少克（【销售结算】）、差多少；'
          '下面每卖一次一行，累计总收入自动加，人民币汇率取【月度汇率】当月的 CNY（1 美元＝多少人民币），折人民币＝金额×汇率。年份空着就跟【首页】。',
          h2=40)
    _sel_row(ws, 'J')
    _lab(ws, 'A3', '年份：', 'B3')
    _inp(ws, 'C3')
    _year_dv(ws, 'C3')
    _note(ws, 'D3', f'="→ "&{Y}&"年"&IF(LEN($C$3&"")=0,"（空着＝跟首页）","")', 'J3')
    ws['L3'] = year_eff('$C$3')
    ws['M3'] = f'=COUNT($M${GD_R0}:$M${R1})'
    for c in 'LMNOP':
        ws[f'{c}3'].font = F_NOTE
    # 一、按月
    ws.merge_cells('A4:J4')
    put(ws, 'A4', '一、按月：冶炼和出售', F_SEC, FILL_SUBH, align=AL, border=False)
    ws.row_dimensions[4].height = 22
    for c, t in zip('ABCDEFG', ('月份', '冶炼纯金量\n（克）', '实际出售纯金量\n（克）', '差额（克）\n冶炼 − 出售', '销售金额\n（美元）',
                                '平均单价\n（美元/克）', '累计销售金额\n（美元）')):
        _hdr(ws, f'{c}5', t, C_SALE)
    ws.row_dimensions[5].height = 34
    pq, pd_ = mod_u(SH_PROD, PD_TON), mod_u(SH_PROD, PD_DATE)
    sq, sd, sa = mod_u(SH_SALE, S_TON), mod_u(SH_SALE, S_DATE), mod_u(SH_SALE, S_AMT)
    for m in range(1, 13):
        r = 5 + m
        d = lambda col: f'{col},">="&DATE({Y},ROW()-5,1),{col},"<"&DATE({Y},ROW()-4,1)'
        put(ws, f'A{r}', f'{m}月', F_TXT, align=AC)
        put(ws, f'B{r}', f'=SUMIFS({pq},{d(pd_)})', F_TXT, fmt=QTY2, align=AR_)
        put(ws, f'C{r}', f'=SUMIFS({sq},{d(sd)})', F_TXT, fmt=QTY2, align=AR_)
        put(ws, f'D{r}', f'=B{r}-C{r}', F_TXT, fmt=QTY2, align=AR_)
        put(ws, f'E{r}', f'=SUMIFS({sa},{d(sd)})', F_TXT, fmt=MONEY, align=AR_)
        put(ws, f'F{r}', f'=IF(N(C{r})=0,"",ROUND(E{r}/C{r},2))', F_TXT, fmt=MONEY, align=AR_)
        put(ws, f'G{r}', f'=N(G{r - 1})+E{r}', F_TXT, fmt=MONEY, align=AR_)
    rt = 18
    put(ws, f'A{rt}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'BCDE':
        put(ws, f'{c}{rt}', f'=SUM({c}6:{c}17)', F_TXTB, FILL_TOT, QTY2 if c != 'E' else MONEY, AR_)
    put(ws, f'F{rt}', f'=IF(N(C{rt})=0,"",ROUND(E{rt}/C{rt},2))', F_TXTB, FILL_TOT, MONEY, AR_)
    put(ws, f'G{rt}', None, F_TXTB, FILL_TOT)
    ws['N3'], ws['O3'] = f'=$C${rt}', f'=$E${rt}'
    # 二、逐笔出售
    ws.merge_cells('A20:J20')
    put(ws, 'A20', '二、逐笔出售记录（取【销售结算】，按出售日期排）', F_SEC, FILL_SUBH, align=AL, border=False)
    ws.row_dimensions[20].height = 22
    for c, t in zip('ABCDEFGHIJ', ('序号', '出售日期', '买方', '实际出售纯金量\n（克）', '单价\n（美元/克）', '实际销售金额\n（美元）',
                                   '总收入\n（累计，美元）', '汇率\n(1美元=X人民币)', '折人民币', '备注')):
        _hdr(ws, f'{c}21', t, C_SALE)
    ws.row_dimensions[21].height = 34
    _helper_head(ws, GD_R0 - 1, [('L', '排序键\n(销售第k行)'), ('M', '名次'), ('N', '第i行是\n销售第几行'), ('P', '折人民币\n(数)')])
    k_ = f'ROW()-{GD_R0 - 1}'
    X = lambda col: f'INDEX({F(SH_SALE, col)},{k_})'
    S = lambda col, r: f'INDEX({F(SH_SALE, col)},$N{r})'
    T = f'{k_}={CNT}+1'
    rtab = f'{q(SH_RATE)}!$B${RATE_R0}:${CL(1 + len(RATE_CURS))}${RATE_R1}'      # 月度汇率 B:I（1 美元＝X 外币），p_wps 换成名称
    rate = lambda r: (f'INDEX({rtab},(YEAR($B{r})-YEAR(起始月份))*12+MONTH($B{r})-MONTH(起始月份)+1,'
                      f'MATCH("{CUR_CNY}",汇率币种头,0))')
    for r in range(GD_R0, R1 + 1):
        ws[f'L{r}'] = (f'=IF(NOT(ISNUMBER({X(S_DATE)})),"",IF(NOT(ISNUMBER({X(S_TON)})),"",IF(YEAR({X(S_DATE)})<>{Y},"",'
                       f'(INT({X(S_DATE)})-DATE({Y},1,1))*1000+{k_})))')
        ws[f'M{r}'] = f'=IF($L{r}="","",COUNTIF($L${GD_R0}:$L${R1},"<="&$L{r}))'
        ws[f'N{r}'] = f'=IF({k_}>{CNT},"",MATCH({k_},$M${GD_R0}:$M${R1},0))'
        ws[f'A{r}'] = f'=IF($N{r}="",IF({T},"合计",""),{k_})'
        ws[f'B{r}'] = f'=IF($N{r}="","",{S(S_DATE, r)})'
        ws[f'C{r}'] = f'=IF($N{r}="","",{S(S_CUS, r)}&"")'
        ws[f'D{r}'] = f'=IF($N{r}="",IF({T},{TG},""),{S(S_TON, r)})'
        ws[f'E{r}'] = f'=IF($N{r}="","",{S(S_PRICE, r)})'
        ws[f'F{r}'] = f'=IF($N{r}="",IF({T},{TU},""),N({S(S_AMT, r)}))'
        ws[f'G{r}'] = f'=IF($N{r}="","",N(G{r - 1})+F{r})'
        ws[f'H{r}'] = f'=IF($N{r}="","",IFERROR(1/(1/{rate(r)}),""))'
        ws[f'P{r}'] = f'=IF($N{r}="","",IF($H{r}="","",ROUND($F{r}*$H{r},2)))'
        ws[f'I{r}'] = f'=IF($N{r}="",IF({T},{TR},""),$P{r})'
        ws[f'J{r}'] = f'=IF($N{r}="","",{S(S_NOTE, r)}&"")'
    ws['P3'] = f'=SUM($P${GD_R0}:$P${R1})'
    _style(ws, 'ABCDEFGHIJ', GD_R0, R1, fmts={'B': DATE, 'D': QTY2, 'E': MONEY0, 'F': MONEY0, 'G': MONEY0, 'H': RATE, 'I': MONEY0},
           aligns={'C': AL, 'J': AL, 'D': AR_, 'E': AR_, 'F': AR_, 'G': AR_, 'H': AR_, 'I': AR_})
    for r in range(GD_R0, R1 + 1):
        for c in 'LMNP':
            ws[f'{c}{r}'].font = F_NOTE
    _cf(ws, f'A{GD_R0}:J{R1}', f'$N{GD_R0}<>""', f'AND($N{GD_R0}="",ROW()-{GD_R0 - 1}={CNT}+1)')
    hide(ws, list('LMNOP'))
    ws.freeze_panes = 'A4'
    _print(ws, f'A1:J{R1}', landscape=True)
    return f'【{SH_GOLD}】按月 12 行＋逐笔出售 {N_SALE} 行'


# ═════════════════════════════ ⑤ 收支分类报表 ═════════════════════════════
# 资金台帐 u4 里她天天看的「收支分类报表」：选年份，收支类别 × 1～12 月（折美元），收入、支出分开（只列这一年有数的类别），带合计。
# 收入＝流水折美元＞0 的、支出＝＜0 的（同一个类别有收有支的两边都列）；没填类别或类别不在【记账规则】里的单列一行。
SZ_N = 170                        # 清单行数：收入、支出各最多 81 行（80 个类别＋没填类别）＋小标题、合计、相抵
SZ_R0 = 5


def build_sz(wb, ctx):
    ws = wb.create_sheet(SH_SZ)
    widths(ws, dict({'A': 24, 'N': 15, 'O': 2}, **{CL(1 + m): 12 for m in range(1, 13)}))
    Y, N1, N2 = '$P$3', '$Q$3', '$R$3'
    R1 = SZ_R0 + SZ_N - 1
    H0, H1 = SZ_R0, SZ_R0 + N_RC              # 隐藏类别表：第 5～84 行＝类别 1～80，第 85 行＝没填类别
    title(ws, f'={Y}&"年收支分类报表（折美元）"', 'N', C_OPS,
          '💡 跟资金台帐里的「收支分类报表」一样：选年份，每个收支类别每个月收了多少、付了多少（都折美元，取【现金流水总表】）。'
          '只列这一年有数的类别；账户之间换汇（购汇/结汇）一边算付、一边算收，「收支相抵」＝三个账户账面美元合计增加了多少。年份空着就跟【首页】。',
          h2=40)
    _sel_row(ws, 'N')
    _lab(ws, 'A3', '年份：')
    _inp(ws, 'B3')
    _year_dv(ws, 'B3')
    _note(ws, 'C3', f'="→ "&{Y}&"年"&IF(LEN($B$3&"")=0,"（空着＝跟首页）","")', 'N3')
    ws['P3'] = year_eff('$B$3')
    ws['Q3'], ws['R3'] = f'=$S${H1}', f'=$T${H1}'
    for c in 'PQR':
        ws[f'{c}3'].font = F_NOTE
    _hdr(ws, 'A4', '收支类别', C_OPS)
    for m in range(1, 13):
        _hdr(ws, f'{CL(1 + m)}4', f'{m}月', C_OPS)
    _hdr(ws, 'N4', '全年合计', C_OPS)
    ws.row_dimensions[4].height = 24
    M, AN, D = cash_u(K_UZS), cash_u(K_CROW), cash_u(K_DATE)
    yr = f'{D},">="&DATE({Y},1,1),{D},"<"&DATE({Y}+1,1,1)'
    _helper_head(ws, SZ_R0 - 1, [('P', '类别'), ('Q', '全年收入'), ('R', '全年支出'), ('S', '收入第几个'), ('T', '支出第几个'),
                                 ('V', '这一行是'), ('W', '第几个类别')])
    k_ = f'ROW()-{H0 - 1}'
    for r in range(H0, H1 + 1):
        if r < H1:
            ws[f'P{r}'] = f'=INDEX(类别表,{k_})&""'
            ws[f'Q{r}'] = f'=IF($P{r}="",0,SUMIFS({M},{AN},{k_},{M},">0",{yr}))'
            ws[f'R{r}'] = f'=IF($P{r}="",0,-SUMIFS({M},{AN},{k_},{M},"<0",{yr}))'
        else:
            ws[f'P{r}'] = '（没填类别 / 类别不在清单里）'
            ws[f'Q{r}'] = f'=SUMIFS({M},{M},">0",{yr})-SUMIFS({M},{AN},">0",{M},">0",{yr})'
            ws[f'R{r}'] = f'=-SUMIFS({M},{M},"<0",{yr})+SUMIFS({M},{AN},">0",{M},"<0",{yr})'
        ws[f'S{r}'] = f'=N(S{r - 1})+IF(ROUND($Q{r},2)<>0,1,0)'
        ws[f'T{r}'] = f'=N(T{r - 1})+IF(ROUND($R{r},2)<>0,1,0)'
    i_ = f'ROW()-{SZ_R0 - 1}'
    for r in range(SZ_R0, R1 + 1):
        V, W = f'$V{r}', f'$W{r}'
        ws[f'V{r}'] = (f'=IF({i_}=1,"H1",IF({i_}<=1+{N1},"I",IF({i_}=2+{N1},"T1",IF({i_}=4+{N1},"H2",'
                       f'IF(AND({i_}>=5+{N1},{i_}<=4+{N1}+{N2}),"E",IF({i_}=5+{N1}+{N2},"T2",IF({i_}=7+{N1}+{N2},"NET","")))))))')
        ws[f'W{r}'] = (f'=IF({V}="I",MATCH({i_}-1,$S${H0}:$S${H1},0),IF({V}="E",MATCH({i_}-4-{N1},$T${H0}:$T${H1},0),""))')
        ws[f'A{r}'] = (f'=IF({V}="H1","一、收入",IF(OR({V}="I",{V}="E"),INDEX($P${H0}:$P${H1},{W}),IF({V}="T1","收入合计",'
                       f'IF({V}="H2","二、支出",IF({V}="T2","支出合计",IF({V}="NET","收支相抵（收入 − 支出）",""))))))')
        for m in range(1, 13):
            dm = f'{D},">="&DATE({Y},{m},1),{D},"<"&DATE({Y},{m + 1},1)'
            inc_t, exp_t = f'SUMIFS({M},{M},">0",{dm})', f'-SUMIFS({M},{M},"<0",{dm})'
            ws[f'{CL(1 + m)}{r}'] = (
                f'=IF({V}="I",IF({W}<={N_RC},SUMIFS({M},{AN},{W},{M},">0",{dm}),{inc_t}-SUMIFS({M},{AN},">0",{M},">0",{dm})),'
                f'IF({V}="E",IF({W}<={N_RC},-SUMIFS({M},{AN},{W},{M},"<0",{dm}),{exp_t}+SUMIFS({M},{AN},">0",{M},"<0",{dm})),'
                f'IF({V}="T1",{inc_t},IF({V}="T2",{exp_t},IF({V}="NET",SUMIFS({M},{dm}),"")))))')
        ws[f'N{r}'] = f'=IF(OR({V}="",{V}="H1",{V}="H2"),"",SUM(B{r}:M{r}))'
    _style(ws, [CL(c) for c in range(1, 15)], SZ_R0, R1, fmts={CL(c): MONEY for c in range(2, 15)},
           aligns=dict({'A': AL}, **{CL(c): AR_ for c in range(2, 15)}))
    for r in range(SZ_R0, max(R1, H1) + 1):
        for c in 'PQRSTVW':
            ws[f'{c}{r}'].font = F_NOTE
    _cf(ws, f'A{SZ_R0}:N{R1}', f'OR($V{SZ_R0}="I",$V{SZ_R0}="E")',
        f'OR($V{SZ_R0}="T1",$V{SZ_R0}="T2",$V{SZ_R0}="NET")', f'OR($V{SZ_R0}="H1",$V{SZ_R0}="H2")')
    hide(ws, list('PQRSTUVW'))
    ws.freeze_panes = f'B{SZ_R0}'
    _print(ws, f'A1:N{R1}', '4:4', landscape=True)
    return f'【{SH_SZ}】收支类别 {N_RC} 个＋没填类别，按年 1～12 月；清单 {SZ_N} 行'


def build(wb, ctx):
    """5 张新报表（建在一本版里；按 plan.md 第三节的顺序，build._final_order 把它们放在【数据校验】前面）"""
    for sh in EXTRA_SHEETS:
        assert sh not in wb.sheetnames, sh
    log = [build_fsq(wb, ctx), build_dg(wb, ctx), build_wz(wb, ctx), build_gold(wb, ctx), build_sz(wb, ctx)]
    assert [ws.title for ws in wb.worksheets if ws.title in EXTRA_SHEETS] == EXTRA_SHEETS
    return log
