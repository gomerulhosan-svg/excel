# -*- coding: utf-8 -*-
"""拆本（split）：在 build 的内存里，从一本版的 wb 出「录入本」「报表本」两本（不读回成品再存：openpyxl 读回再存会丢外部链接关系、
丢合并区里藏着的公式）。用法见 build.build_split：

    undo = split.make_report(wb, ctx)   →  wb.save(报表本)  →  undo()      （报表本的改动全记下来，存完原样退回）
    split.make_entry(wb, ctx)           →  wb.save(录入本)                  （录入本是拆掉的，最后做）

报表本（make_report）：一本版原样（报表、隐藏计算表、p_wps 的分段名称一个字不改），只把录入本那些表变成同名、同位置的隐藏取数表：
  · common.SPLIT 里「填」的格子（录入表的填列、基础资料等的填块；录入表表头行的填列也取，数据校验能比「录入本表头挪没挪」）
    换成取数公式 mirror_formula()：MIRROR_MODE='cell'（默认，写法①）每格 =IFERROR(IF([1]表!A6="","",[1]表!A6),"")；
    'name'（写法③）每个填列一个隐藏名称 录_流水_A＝[1]现金流水总表!$A$6:$A$3965，每格 INDEX(名称,ROW()-5)（能压共享）；
  · result 表（录入检查）：只取 entry.ECHK_MIRROR 那些格子（汇总、各项结果），报表本直接看录入本算出来的检查结果；
    一本版里【录入检查】由 build.make_onebook 调 entry.build_entry_check 建好（报表本【数据校验】第 41 项取它）；
  · whole 表（货柜登记）：报表本里没人引用就不要这张表（引用了就整张取数）；
  · 取数表隐藏，去掉数据有效性和条件格式；外部链接（openpyxl ExternalLink，目标＝相对文件名 WB1_FILE）这次存盘一起建。
录入本（make_entry，顺序照 entry.py 的说明）：
  · 只留 WB1_ORDER 的表（报表、_凭证索引、_结转损益、旧首页删掉）；
  · entry.apply(wb, ctx)：录入本写法（ENTRY_FORMULAS）、清空重列（ENTRY_CLEAR）、只在报表本算的列刷灰加注、_辅助 只留末行和块前累计；
    SPLIT 的「重」列这里再兜底清一遍；_辅助 再按引用一层层找，用不上的公式格清掉；
  · entry.build_entry_home（新首页）、entry.build_entry_help（有就调；没有就留一本版的使用说明）、entry.build_entry_check；
  · 名称只留用得上的（公式、数据有效性、条件格式用到的，连名称套名称）；
  · 留下的公式要是还引用删掉的表 / 只有报表才有的名称 / 清空了的重列，就报错（STRICT=False 时只记下来、把那格清掉，测试用）。
"""
import importlib
import re
from collections import defaultdict

from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidationList
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.cell.cell import MergedCell

import common
from common import (SPLIT, WB1_FILE, WB1_ORDER, EXT_PREFIX, SH_HOME, SH_HELP, SH_AUX, SEG, FILL_AUTO, q)

import importlib.util
# core 的 entry.py（录入本写法、首页、录入检查）；还没有时也能出两本（首页占位）。entry 里自己的错照常报出来
entry = importlib.import_module('entry') if importlib.util.find_spec('entry') else None

STRICT = True                # False：留下的公式引用了删掉的东西，只记下来并清掉那格（测试用），不报错
NOTE = '（在报表本看）'

# ─────────────────────────── 公式里的引用 ───────────────────────────
_ID = r'[A-Za-z_一-鿿][\w一-鿿.]*'
QREF = re.compile(r"(?:'((?:[^']|'')+)'|(" + _ID + r"))!(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?(?![\w(])")
LREF = re.compile(r"(?<![\w.!$'\]])(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?(?![\w(!])")
TOK = re.compile(_ID)
CELLP = re.compile(r'\$?([A-Z]{1,3})\$?(\d+)$')


def _is_f(v):
    return isinstance(v, str) and v.startswith('=')


def _cols(s):
    return s.split() if isinstance(s, str) else list(s)


def _cell(a):
    m = CELLP.match(a)
    return CI(m.group(1)), int(m.group(2))


def _box(a, b):
    c0, r0 = _cell(a)
    c1, r1 = _cell(b) if b else (c0, r0)
    return min(c0, c1), min(r0, r1), max(c0, c1), max(r0, r1)


def refs(f):
    """公式（或名称、有效性公式）→ (带表名的区域 [(表, c0, r0, c1, r1)], 本表区域 [(c0, r0, c1, r1)], 标识符集合)。字符串里的不算"""
    body = ''.join(f.split('"')[0::2])
    qs = []

    def qsub(m):
        sh = (m.group(1) or '').replace("''", "'") if m.group(1) is not None else m.group(2)
        qs.append((sh,) + _box(m.group(3), m.group(4)))
        return ' '
    rest = QREF.sub(qsub, body)
    ls = []

    def lsub(m):
        ls.append(_box(m.group(1), m.group(2)))
        return ' '
    rest = LREF.sub(lsub, rest)
    return qs, ls, set(TOK.findall(rest))


def _hit(box, area):
    c0, r0, c1, r1 = box
    a0, b0, a1, b1 = area
    return not (c1 < a0 or c0 > a1 or r1 < b0 or r0 > b1)


def _formulas(ws):
    for row in ws.iter_rows():
        for c in row:
            if _is_f(c.value):
                yield c


def _dv_cf_formulas(ws):
    """数据有效性、条件格式里的公式（前面补「=」）"""
    out = []
    for dv in ws.data_validations.dataValidation:
        for f in (dv.formula1, dv.formula2):
            if f:
                out.append('=' + f)
    for cf in ws.conditional_formatting:
        for rule in cf.rules:
            for f in (rule.formula or []):
                out.append('=' + f)
    return out


# ─────────────────────────── 拆本规格（common.SPLIT）落到这本上 ───────────────────────────
def check_spec(wb, log=print):
    """SPLIT → 每张表实际要取数 / 清空的格子；有列没归类就报错。
    返回 {表: dict(kind, fill=[(列, 起行, 止行)], hdr=[格子], clear=[(列, 起行, 止行)])}"""
    out = {}
    bad = []
    for sh, sp in SPLIT.items():
        if sh not in wb.sheetnames:
            log(f'   · 拆本：一本版里没有【{sh}】，跳过')
            continue
        ws = wb[sh]
        if sp.get('whole') or sp.get('result'):
            out[sh] = dict(kind='whole' if sp.get('whole') else 'result', fill=[], hdr=[], clear=[])
            continue
        hdr = sp['hdr']
        if 'blocks' in sp:
            blocks = list(sp['blocks'])
            for c, r0, r1 in blocks:
                for r in range(max(r0, hdr + 1), r1 + 1):
                    if _is_f(ws[f'{c}{r}'].value):
                        bad.append(f'{sh}!{c}{r} 在「填」块里却是公式')
                        break
            if sp.get('auto'):
                # 表头行以下整块（所有列、到最后一行）凡不是公式的格子都取数（make_report 跳过公式格）：
                # 二级明细对照这种新块不管放在右边还是下面都接得上，不用在这里一块块登记
                known = {(CI(c), r) for c, r0, r1 in blocks for r in range(r0, r1 + 1)}
                more = sum(1 for row in ws.iter_rows(min_row=hdr + 1) for c in row
                           if c.value is not None and not _is_f(c.value) and (c.column, c.row) not in known)
                blocks += [(CL(ci), hdr + 1, ws.max_row) for ci in range(1, ws.max_column + 1)]
                if more:
                    log(f'   · 拆本：【{sh}】明列的块以外还有 {more} 个有字的格子（新块），第 {hdr + 1}～{ws.max_row} 行整块取数')
            else:
                covered = {c for c, _, _ in blocks} | set(_cols(sp.get('文字', '')))
                fcols, hcols = set(), {}
                for row in ws.iter_rows(min_row=hdr):
                    for c in row:
                        if c.value is None:
                            continue
                        if c.row == hdr:
                            hcols[c.column_letter] = c.value
                        elif _is_f(c.value):
                            fcols.add(c.column_letter)
                for L in sorted((L for L in hcols if L not in covered and L not in fcols), key=CI):
                    bad.append(f'{sh}!{L}{hdr}「{str(hcols[L])[:16]}」没归类（不是公式列、也不在 SPLIT 的填块里）')
            out[sh] = dict(kind='blocks', fill=blocks, hdr=[], clear=[])
            continue
        r0, r1 = sp['r0'], sp['r1']
        fill, light, heavy = _cols(sp['填']), _cols(sp['轻']), _cols(sp['重'])
        allc = fill + light + heavy
        dup = {c for c in allc if allc.count(c) > 1}
        if dup:
            bad.append(f'{sh} 列 {" ".join(sorted(dup, key=CI))} 归了两类')
        seen = {}
        for row in ws.iter_rows(min_row=hdr, max_row=r1):
            for c in row:
                if c.value is not None:
                    seen.setdefault(c.column_letter, c)
        for L, c in sorted(seen.items(), key=lambda x: CI(x[0])):
            if L not in allc:
                hv = ws[f'{L}{hdr}'].value
                bad.append(f'{sh}!{L} 列（表头「{str(hv or "")[:16]}」）没归进 填/轻/重')
        for L in fill:
            for r in range(r0, r1 + 1):
                if _is_f(ws[f'{L}{r}'].value):
                    bad.append(f'{sh}!{L}{r} 是「填」列却是公式')
                    break
        out[sh] = dict(kind='table', fill=[(c, r0, r1) for c in fill], hdr=[f'{c}{hdr}' for c in fill if ws[f'{c}{hdr}'].value is not None],
                       clear=[(c, r0, r1) for c in heavy])
    if bad:
        raise AssertionError('拆本规格（common.SPLIT）跟表对不上：\n  ' + '\n  '.join(bad[:40]))
    return out


# ─────────────────────────── 报表本 ───────────────────────────
def _xref(sh, a):
    """[1]表!A6（表名要引号时 '[1]表'!A6）"""
    return (f"'[1]{sh}'!{a}" if q(sh).startswith("'") else f'[1]{sh}!{a}')


def _wrap(x):
    f = f'IF({x}="","",{x})'
    return f'=IFERROR({f},"")' if common.MIRROR_IFERROR else '=' + f


def mirror_formula(sh, col, r, name=None, r0=None):
    """取数公式。name/r0 给了（MIRROR_MODE='name'）＝INDEX(名称,ROW()-常数)；否则每格直接链接"""
    if name:
        return _wrap(f'INDEX({name},ROW()-{r0 - 1})')
    return _wrap(_xref(sh, f'{col}{r}'))


def _ext_name(sh, col, r0):
    p = SEG[sh][0] if sh in SEG else sh
    return f'{EXT_PREFIX}{p}_{col}' + ('' if sh in SEG else f'{r0}')


def add_external_link(wb, target, sheet_names):
    """报表本 → 录入本的外部链接（照 A065 s_mirror.add_external_link，目标写相对文件名）。返回链接对象（undo 用）"""
    from openpyxl.workbook.external_link.external import (ExternalLink, ExternalBook, ExternalSheetNames,
                                                           ExternalSheetDataSet, ExternalSheetData)
    from openpyxl.packaging.relationship import Relationship
    eb = ExternalBook(sheetNames=ExternalSheetNames(sheetName=list(sheet_names)),
                      sheetDataSet=ExternalSheetDataSet(sheetData=[ExternalSheetData(sheetId=i) for i in range(len(sheet_names))]),
                      id='rId1')
    el = ExternalLink(externalBook=eb)
    el.file_link = Relationship(type='externalLinkPath', Target=target, TargetMode='External', Id='rId1')
    assert not wb._external_links, '一本版里已经有外部链接了'
    wb._external_links.append(el)
    return el


def _referenced(wb, sh):
    """sh 有没有被别的表的公式 / 有效性 / 条件格式引用（直接写表名，或者通过指向它的名称）"""
    nms = [n for n, dn in wb.defined_names.items() if sh in (dn.attr_text or '')]
    for ws in wb.worksheets:
        if ws.title == sh:
            continue
        fs = [c.value for c in _formulas(ws)] + _dv_cf_formulas(ws)
        for f in fs:
            f = ''.join(f.split('"')[0::2])          # 字符串里的表名不算
            if sh in f:
                return True
            if nms and any(n in f for n in nms):
                return True
    return False


def make_report(wb, ctx, log=print):
    """一本版 → 报表本（就地改，返回 undo()：调了就原样退回一本版）"""
    spec = check_spec(wb, log)
    mode = common.MIRROR_MODE
    assert mode in ('cell', 'name'), mode
    olds = []                    # (ws, 格子, 原值)
    states = []                  # (ws, 原 sheet_state)
    dvcf = []                    # (ws, 原有效性, 原条件格式)
    new_names = []
    removed = []                 # (位置, ws, {名称: DefinedName})
    n_cell = defaultdict(int)

    def setv(ws, a, v):
        olds.append((ws, a, ws[a].value))
        ws[a].value = v

    def ext_name(sh, col, r0, r1):
        nm = _ext_name(sh, col, r0)
        if nm not in wb.defined_names:
            wb.defined_names[nm] = DefinedName(nm, attr_text=_xref(sh, f'${col}${r0}:${col}${r1}'), hidden=True)
            new_names.append(nm)
        return nm

    for sh, s in spec.items():
        ws = wb[sh]
        if s['kind'] == 'whole' and not _referenced(wb, sh):
            idx = wb._sheets.index(ws)
            gone = {n: dn for n, dn in wb.defined_names.items() if sh in (dn.attr_text or '')}
            for n in gone:
                del wb.defined_names[n]
            wb._sheets.remove(ws)
            removed.append((idx, ws, gone))
            log(f'   · 报表本：【{sh}】没有公式引用，不放进报表本')
            continue
        if s['kind'] == 'result' and _entry_attr('ECHK') == sh and _entry_attr('ECHK_MIRROR'):
            # 录入检查：只取 entry.ECHK_MIRROR 那些格子（汇总、各项结果）；别的公式（辅助列）报表本里用不着，清掉
            want = set()
            for x in entry.ECHK_MIRROR:
                c0, r0, c1, r1 = _box(*x.split(':')) if ':' in x else _box(x, None)
                want |= {f'{CL(c)}{r}' for c in range(c0, c1 + 1) for r in range(r0, r1 + 1)}
            for c in list(_formulas(ws)):
                if c.coordinate not in want:
                    setv(ws, c.coordinate, None)
            for a in sorted(want):
                setv(ws, a, mirror_formula(sh, *_split_a(a)))
                n_cell[sh] += 1
        elif s['kind'] in ('whole', 'result'):
            for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
                for c in row:
                    if c.value is None and s['kind'] == 'result':
                        continue
                    if isinstance(c, MergedCell) or (s['kind'] == 'whole' and _is_f(c.value)):
                        continue
                    setv(ws, c.coordinate, mirror_formula(sh, c.column_letter, c.row))
                    n_cell[sh] += 1
        else:
            for col, r0, r1 in s['fill']:
                nm = ext_name(sh, col, r0, r1) if mode == 'name' else None
                for r in range(r0, r1 + 1):
                    a = f'{col}{r}'
                    if _is_f(ws[a].value) or isinstance(ws[a], MergedCell):
                        continue
                    setv(ws, a, mirror_formula(sh, col, r, nm, r0))
                    n_cell[sh] += 1
            for a in s['hdr']:                  # 表头（录入本表头挪了，报表本数据校验能看出来）
                if not _is_f(ws[a].value) and not isinstance(ws[a], MergedCell):
                    setv(ws, a, mirror_formula(sh, *_split_a(a)))
                    n_cell[sh] += 1
        states.append((ws, ws.sheet_state))
        ws.sheet_state = 'hidden'
        dvcf.append((ws, ws.data_validations, ws.conditional_formatting))
        ws.data_validations = DataValidationList()
        ws.conditional_formatting = ConditionalFormattingList()
    link = add_external_link(wb, WB1_FILE, list(WB1_ORDER))      # 表名清单以录入本实际的为准（refresh_link_cache 按录入本重写）
    act = wb.active
    log(f'   · 报表本：取数表 {len(n_cell)} 张、取数公式 {sum(n_cell.values()):,} 格（写法 {mode}'
        + (f'，外部名称 {len(new_names)} 个' if new_names else '') + '）：'
        + '、'.join(f'{k} {v:,}' for k, v in n_cell.items()))

    def undo():
        wb._external_links.remove(link)
        for ws, a, v in reversed(olds):
            ws[a].value = v
        for ws, st in states:
            ws.sheet_state = st
        for ws, dv, cf in dvcf:
            ws.data_validations = dv
            ws.conditional_formatting = cf
        for nm in new_names:
            del wb.defined_names[nm]
        for idx, ws, gone in reversed(removed):
            wb._sheets.insert(idx, ws)
            for n, dn in gone.items():
                wb.defined_names[n] = dn
        wb.active = act
    return undo


def _split_a(a):
    m = CELLP.match(a)
    return m.group(1), int(m.group(2))


# ─────────────────────────── 录入本 ───────────────────────────
def _entry_attr(name, default=None):
    return getattr(entry, name, default) if entry is not None else default


def make_entry(wb, ctx, log=print):
    """一本版（已经做完报表本、undo 过）→ 录入本（就地拆）。顺序照 entry.py 的说明：删报表表 → entry.apply → 首页 → 录入检查 → 过滤名称"""
    spec = check_spec(wb, lambda *a: None)
    keep = {s for s in WB1_ORDER if s in wb.sheetnames and s != SH_HOME}
    gone = [ws.title for ws in wb.worksheets if ws.title not in keep]
    names = {n: dn.attr_text or '' for n, dn in wb.defined_names.items()}

    # 1) 删报表表（名称先留着，查完再删）；entry.apply：录入本写法、清空重列、_辅助 只留末行和块前累计
    for sh in gone:
        del wb[sh]
    if _entry_attr('apply'):
        for line in entry.apply(wb, ctx):
            log('   · 录入本 ' + line)
    else:
        log('   ⚠ 没有 entry.py：录入本只清空重列，不换写法')
    cleared = defaultdict(list)          # 表 → [(c0, r0, c1, r1)]：SPLIT 的重列（entry.ENTRY_CLEAR 没清到的这里兜底）
    n_clr = 0
    for sh, s in spec.items():
        if sh not in keep:
            continue
        ws = wb[sh]
        for c, r0, r1 in s['clear']:
            cleared[sh].append((CI(c), r0, CI(c), r1))
            vis = entry is None and not ws.column_dimensions[c].hidden
            for r in range(r0, r1 + 1):
                cell = ws[f'{c}{r}']
                if cell.value is not None:
                    cell.value = None
                    n_clr += 1
                if vis:
                    cell.fill = FILL_AUTO
            h = ws[f'{c}{SPLIT[sh]["hdr"]}']
            if vis and isinstance(h.value, str) and '报表本' not in h.value:
                h.value = h.value + '\n' + NOTE
    if n_clr:
        log(f'   · 录入本：SPLIT 重列兜底又清了 {n_clr:,} 格（entry.ENTRY_CLEAR 没清到的）')

    # 2) 哪些名称不能用了：指向删掉的表，或者落在清空的重列里（名称套名称一起算）
    dead = {}
    for n, t in names.items():
        for sh, *box in refs('=' + t)[0]:
            if sh in gone:
                dead[n] = f'指向【{sh}】'
            elif sh in cleared and any(_hit(box, ar) for ar in cleared[sh]):
                dead[n] = f'落在【{sh}】清空的列'
    ch = True
    while ch:
        ch = False
        for n, t in names.items():
            if n in dead:
                continue
            _, _, toks = refs('=' + t)
            hit = toks & set(dead)
            if hit:
                dead[n] = f'用到名称 {sorted(hit)[0]}'
                ch = True

    # 3) _辅助 只留用得上的格子：从留下的表的公式 / 有效性 / 条件格式出发，顺着名称和引用一层层找
    aux = wb[SH_AUX] if SH_AUX in keep else None
    need = set()
    used_names = set()
    todo_f = []
    for sh in keep:
        if sh == SH_AUX:
            continue
        ws = wb[sh]
        todo_f += [(sh, c.coordinate, c.value) for c in _formulas(ws)]
        todo_f += [(sh, None, f) for f in _dv_cf_formulas(ws)]
    offenders = defaultdict(list)

    def scan(sh, a, f):
        """记下这条公式用到的名称和 _辅助 格子；引用了删掉 / 清空的东西就记成问题"""
        qs, ls, toks = refs(f)
        new_aux = []
        bad = []
        for s2, *box in qs:
            if s2 == SH_AUX:
                new_aux.append(box)
            elif s2 in gone:
                bad.append(f'{s2}!')
            elif s2 in cleared and any(_hit(box, ar) for ar in cleared[s2]):
                bad.append(f'{s2}!{CL(box[0])}（清空的列）')
        for box in ls:
            if sh == SH_AUX:
                new_aux.append(box)
            elif sh in cleared and any(_hit(box, ar) for ar in cleared[sh]):
                bad.append(f'{CL(box[0])}{box[1]}（清空的列）')
        for t in toks & set(names):
            if t in dead:
                bad.append(f'{t}（{dead[t]}）')
            elif t not in used_names:
                used_names.add(t)
                qn, _, tn = refs('=' + names[t])
                new_aux += [box for s2, *box in qn if s2 == SH_AUX]
                for t2 in tn & set(names):
                    if t2 not in used_names and t2 not in dead:
                        todo_names.append(t2)
        if bad and a is not None:
            col = re.match(r'[A-Z]+', a).group(0)
            offenders[(sh, col)].append((a, bad[0]))
        elif bad:
            offenders[(sh, '有效性/条件格式')].append(('', bad[0]))
        return new_aux

    todo_names = []
    boxes = []
    for sh, a, f in todo_f:
        boxes += scan(sh, a, f)
    while todo_names:
        t = todo_names.pop()
        if t in used_names:
            continue
        boxes += scan('', None, '=' + t)
    if aux is not None:
        by_row = defaultdict(list)            # 行 → [(列号, 格子)]（_辅助 的公式格）
        for c in _formulas(aux):
            by_row[c.row].append((c.column, c))
        while boxes:
            c0, r0, c1, r1 = boxes.pop()
            for rr in range(r0, min(r1, aux.max_row) + 1):
                for cc, cell in by_row.get(rr, ()):
                    if c0 <= cc <= c1 and (cc, rr) not in need:
                        need.add((cc, rr))
                        boxes += scan(SH_AUX, cell.coordinate, cell.value)
                        while todo_names:
                            t = todo_names.pop()
                            if t not in used_names:
                                boxes += scan('', None, '=' + t)
    if offenders:
        lines = [f'{sh}!{col}：{len(v)} 格，比如 {v[0][0]} 用到 {v[0][1]}' for (sh, col), v in sorted(offenders.items())]
        msg = '录入本里留下的公式还引用了报表本才有的东西（在 entry.ENTRY_FORMULAS / ENTRY_CLEAR 里给录入本写法）：\n  ' + '\n  '.join(lines)
        if STRICT:
            raise AssertionError(msg)
        log('   ⚠ ' + msg.replace('\n', '\n   '))
        for (sh, col), v in offenders.items():
            if col != '有效性/条件格式':
                for a, _ in v:
                    wb[sh][a].value = None

    # 4) 删名称、_辅助 剪枝、去掉指向删掉的表的链接
    if aux is not None:
        rows = {r for _, r in need}
        n_drop = 0
        for row in aux.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                if _is_f(c.value):
                    if (c.column, c.row) not in need:
                        c.value = None
                        n_drop += 1
                elif not ({c.row, c.row + 1, c.row + 2} & rows):
                    c.value = None
        log(f'   · 录入本：_辅助 留 {len(need)} 个公式格（第 {min(rows) if rows else 0}～{max(rows) if rows else 0} 行里的），清掉 {n_drop:,} 个')
    n_link = 0
    for sh in keep:
        for row in wb[sh].iter_rows():
            for c in row:
                h = c.hyperlink
                if h is not None and h.location:
                    m = re.match(r"^'?((?:[^']|'')+?)'?!", h.location)
                    if m and m.group(1) in gone and m.group(1) not in (SH_HOME, SH_HELP):
                        c.hyperlink = None
                        n_link += 1
    if n_link:
        log(f'   · 录入本：去掉指向报表的链接 {n_link} 个')
    for n in list(wb.defined_names):
        if n in dead:
            del wb.defined_names[n]

    # 5) 新首页、使用说明、录入检查
    hb = _entry_attr('build_entry_home')
    if hb:
        hb(wb, ctx)
    else:
        _placeholder(wb, SH_HOME, '录入本 · 首页（待 entry.build_entry_home）')
    hh = _entry_attr('build_entry_help')
    if hh:
        hh(wb, ctx)
    else:
        _placeholder(wb, SH_HELP, '录入本 · 使用说明（待 entry.build_entry_help）')
    if _entry_attr('build_entry_check'):
        entry.build_entry_check(wb, ctx)           # 一本版里那张换掉（同一个函数建的，一样）

    # 6) 名称只留用得上的
    final = {n: dn.attr_text or '' for n, dn in wb.defined_names.items()}
    used = set()
    stack = []
    for ws in wb.worksheets:
        for f in [c.value for c in _formulas(ws)] + _dv_cf_formulas(ws):
            stack += list(refs(f)[2] & set(final))
    while stack:
        t = stack.pop()
        if t in used:
            continue
        used.add(t)
        stack += list(refs('=' + final[t])[2] & set(final))
    n0 = len(final)
    for n in list(final):
        if n not in used:
            del wb.defined_names[n]
    missing = sorted({t for ws in wb.worksheets for c in _formulas(ws) for t in refs(c.value)[2] if t in names and t not in final})
    if missing:
        msg = f'录入本的公式用到了已经删掉的名称：{missing[:20]}'
        if STRICT:
            raise AssertionError(msg)
        log('   ⚠ ' + msg)
    log(f'   · 录入本：名称 {len(names):,} → {len(used):,} 个（删掉指向报表的 {len(dead):,} 个、用不上的 {n0 - len(used):,} 个）')

    # 7) 表顺序、隐藏、活动表
    order = [s for s in WB1_ORDER if s in wb.sheetnames]
    order = [s for s in order if s != SH_AUX] + [s for s in wb.sheetnames if s not in order] + ([SH_AUX] if SH_AUX in wb.sheetnames else [])
    wb._sheets = [wb[s] for s in order]
    if SH_AUX in wb.sheetnames:
        wb[SH_AUX].sheet_state = 'hidden'
    wb.active = 0
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == order[0]
    log(f'   · 录入本：{len(wb.sheetnames)} 张表：{"、".join(wb.sheetnames)}')


def _placeholder(wb, sh, text):
    if sh in wb.sheetnames:
        return
    ws = wb.create_sheet(sh)
    ws['A1'] = text
