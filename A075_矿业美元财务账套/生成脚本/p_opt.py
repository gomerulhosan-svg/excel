# -*- coding: utf-8 -*-
"""提速（opt）里要在 p1008 之后做的几处：
  · 凭证汇总、记账凭证按「凭证所在的那一块」查分录（apply 前半）；
  · 数据校验第 41 项「没有插入 / 删除整行」（step_rows_check）；
  · 分录精简：有凭证的分录挤到隐藏表 _分录精简，科目余额表 / 往来余额表 / _结转损益 / 经营报表走势只扫这几行（step_compact）；
  · _凭证索引「当月第几号」COUNTIFS → COUNTIF 减本月以前的张数（step_vx_rank）。

下面是第一项的说明。

记账分录是按来源分块排的（流水 / 采购 / 领用 / 销售 / 工资 / 手工 / 折旧 / 调汇 / 成本 / 销售成本 / 兑换 / 结转损益 / 年末结转，
各块的起止行见 s_je.layout 和 p1008），一张凭证的分录只会在它自己那一块里、从「凭证ID 对应的那个槽位」往下：
  · 凭证ID＝来源代码×100000＋位置（结转损益 12、年末结转 13），位置＝这张凭证第一行在源表里是第几行（月末几种＝月份序号）；
  · 合并成一张的几笔（按天合并、合并号、单据号、手工凭证组号），第一笔就是凭证ID 指的那一笔，其余都在它下面。
所以原来「SUMIF(分录凭证ID,…)」「MATCH(查找键,分录键,0)」扫整张记账分录 1.8 万行的，改成只扫这一段（INDEX(…,起):INDEX(…,止)），
结果一样：凭证汇总 E/F 借贷合计、记账凭证每条分录在第几行、合计行。
凭证汇总 N 列（索引行）：超过本月凭证张数的行不再 MATCH（原来 450 行每行都在 6,898 行里找一遍，找不到才停）。
"""
from common import *
import p1008
from s_je import VL, VL_R1, VP


def windows(ctx):
    """[(相对起行, 每个位置几行, 块末相对行)] 按来源代码 1..13（相对行＝记账分录第 4 行算第 1 行）"""
    bl = ctx['blocks']
    out = {}
    for src, b in bl.items():
        o = b['je0'] - JE_R0 + 1
        out[SRC[src]] = (o, b['legs'], o + b['n'] * b['legs'] - 1)
    o = p1008.CLOSE0 - JE_R0 + 1
    out[12] = (o, p1008.MBLK, o + p1008.NM * p1008.MBLK - 1)
    o = p1008.YE0 - JE_R0 + 1
    out[13] = (o, 2, o + p1008.NY * 2 - 1)
    assert sorted(out) == list(range(1, 14)), sorted(out)
    assert out[13][2] == ctx['je_end'] - JE_R0 + 1
    return [out[i] for i in range(1, 14)]


def win_formulas(vid, w):
    s = f'INT({vid}/100000)'
    start = (f'CHOOSE({s},{",".join(str(x[0]) for x in w)})+(MOD({vid},100000)-1)*'
             f'CHOOSE({s},{",".join(str(x[1]) for x in w)})')
    end = f'CHOOSE({s},{",".join(str(x[2]) for x in w)})'
    return start, end


def seg(name, a, b):
    return f'INDEX({name},{a}):INDEX({name},{b})'


def apply(wb, ctx):
    log = []
    w = windows(ctx)
    # ── 凭证汇总 ──
    ws = wb[SH_VLIST]
    c = VL
    put(ws, 'N4', '=COUNTIFS(凭证月份,报表年度*100+报表月份,凭证有,1)', F_NOTE, border=False)
    put(ws, 'M4', '本月张数→', F_NOTE, border=False)
    mk = '(报表年度*100+报表月份)'
    for k in range(1, VL_N + 1):
        r = VL_R0 + k - 1
        X, V = f'${c["idx"]}{r}', f'${c["vid"]}{r}'
        old = ws[f'{c["idx"]}{r}'].value
        assert old == f'=IFERROR(MATCH({mk}*10000+{k},凭证查找键,0),"")', old
        ws[f'{c["idx"]}{r}'] = f'=IF({k}>$N$4,"",IFERROR(MATCH({mk}*10000+{k},凭证查找键,0),""))'
        st, en = win_formulas(V, w)
        ws[f'P{r}'] = f'=IF({V}="","",{st})'
        ws[f'Q{r}'] = f'=IF({V}="","",{en})'
        for col, side in ((c['dr'], '分录借方'), (c['cr'], '分录贷方')):
            assert ws[f'{col}{r}'].value == f'=IF({X}="","",SUMIF(分录凭证ID,{V},{side}))'
            ws[f'{col}{r}'] = (f'=IF({X}="","",SUMIF({seg("分录凭证ID", f"$P{r}", f"$Q{r}")},{V},'
                               f'{seg(side, f"$P{r}", f"$Q{r}")}))')
        for col in ('P', 'Q'):
            ws[f'{col}{r}'].font = F_NOTE
    for col, t in (('P', '分录起'), ('Q', '分录止')):
        put(ws, f'{col}{VL_HDR}', t, F_NOTE, FILL_AUTO, align=AC)
        ws.column_dimensions[col].hidden = True
    ws.column_dimensions['M'].hidden = True       # M、N 本来就是隐藏的辅助列（凭证ID、索引行）
    ws.column_dimensions['N'].hidden = True
    log.append(f'凭证汇总：索引行超过本月张数不再查；借贷合计只扫这张凭证所在的那一块（P、Q 隐藏列＝分录起止行）')

    # ── 记账凭证 ──
    ws = wb[SH_VPRT]
    v = VP
    n = 0
    for p in range(1, VP_PAGES + 1):
        R = (p - 1) * VP_ROWS + 1
        P0, VID = f'${v["page"]}${R}', f'${v["vid"]}${R}'
        st, en = win_formulas(VID, w)
        ws[f'{v["key"]}{R}'] = f'=IF({P0}="","",{st})'
        ws[f'{v["jrow"]}{R}'] = f'=IF({P0}="","",{en})'
        S, E = f'${v["key"]}${R}', f'${v["jrow"]}${R}'
        for c_ in (v['key'], v['jrow']):
            ws[f'{c_}{R}'].font = F_NOTE
        for i in range(1, VP_LINES + 1):
            r = R + 2 + i
            key = f'${v["key"]}{r}'
            old = ws[f'{v["jrow"]}{r}'].value
            assert old == f'=IF({P0}="","",IFERROR(MATCH({key},分录键,0),""))', old
            ws[f'{v["jrow"]}{r}'] = f'=IF({P0}="","",IFERROR(MATCH({key},{seg("分录键", S, E)},0)+{S}-1,""))'
            n += 1
        r = R + 3 + VP_LINES
        for col, side in ((v['dr'], '分录借方'), (v['cr'], '分录贷方')):
            old = ws[f'{col}{r}'].value
            assert old.startswith('=IF(IF(') and f'SUMIF(分录凭证ID,{VID},{side})' in old, old
            ws[f'{col}{r}'].value = old.replace(f'SUMIF(分录凭证ID,{VID},{side})',
                                                f'SUMIF({seg("分录凭证ID", S, E)},{VID},{seg(side, S, E)})')
    log.append(f'记账凭证：每条分录（{n} 格）、合计行只在这张凭证所在的那一块里找（P、Q 两列第一行＝分录起止行）')
    log += step_rows_check(wb)
    log += step_compact(wb, ctx)
    log += step_vx_rank(wb, ctx)
    return log


# ─────────────────────────── 分录精简 ───────────────────────────
# 记账分录一共 18,018 个槽位（流水 4000×2、各模块、月末、结转损益 32×61……），真正有凭证的行通常只有一两成。
# 科目余额表、往来余额表、_结转损益、经营报表走势原来每格都把 1.8 万行扫一遍；这里把「有凭证的行」按原来的先后顺序
# 挤到一张隐藏表里（只挤这些公式要用的 7 列），这些公式改成只扫「挤过来的行数」那么长。
# 结果一样：没凭证的行（L 列凭证ID 是空的）日期、科目、余额键、分类键、往来、损益键都是空的，本来就不会被算进去；
# 挤过来的行顺序没变，SUMIF 加的还是同一串数、同一个先后顺序。
# ⚠ 结转损益块、年末结转块（第 CLOSE0 行起）不管有没有凭证都整块挤过来：这两块的「有没有凭证」取决于 _结转损益 算出来的数，
#   而 _结转损益 要用精简表的损益键——要是按「有没有凭证」挤，就成了循环引用。整块挤过来不影响结果（空行的键都是空的），只是多扫 1,958 行。
CMP_R0 = 4
CMP_COLS = [('B', J_TBKEY, '余额键', '精简余额键'), ('C', J_DR, '借方', '精简借方'),
            ('D', J_CR, '贷方', '精简贷方'), ('E', J_CLSKEY, '分类键', '精简分类键'), ('F', J_CP, '往来单位', '精简往来'),
            ('G', 'R', '损益键', '精简损益键'), ('H', 'S', '损益额', '精简损益额')]
J_CNT = 'X'                     # 记账分录 X 列：到本行为止挤进精简表的行数


def _swap(ws, coord, pairs, must=True):
    v = ws[coord].value
    w = v
    for a, b in pairs:
        w = w.replace(a, b)
    if must:
        assert w != v, (ws.title, coord, v[:200])
    ws[coord].value = w
    return w


def step_compact(wb, ctx):
    from openpyxl.workbook.defined_name import DefinedName
    log = []
    R0, R1 = JE_R0, ctx['je_end']
    N = R1 - R0 + 1
    JE = q(SH_JE)
    je = wb[SH_JE]
    assert je[f'{J_CNT}{R0}'].value is None and je[f'{J_CNT}{R1}'].value is None
    put(je, f'{J_CNT}3', '挤进精简表\n累计行数', F_NOTE, FILL_AUTO, align=ACW)
    for r in range(R0, R1 + 1):
        prev = '' if r == R0 else f'${J_CNT}{r - 1}+'
        je[f'{J_CNT}{r}'] = f'={prev}IF(${J_VID}{r}="",0,1)' if r < p1008.CLOSE0 else f'={prev}1'
        je[f'{J_CNT}{r}'].font = F_NOTE
    je.column_dimensions[J_CNT].hidden = True

    ws = wb.create_sheet(SH_CMP)
    ws.sheet_state = 'hidden'
    ws['A1'] = (f'分录精简（自动，不用动）：【记账分录】里有凭证的行（L 列凭证ID 不空；第 {p1008.CLOSE0} 行起的结转损益、年末结转块整块），'
                '按原来的先后顺序挤到这里；科目余额表、往来余额表、_结转损益、经营报表走势只扫这几行（名称「精简…」＝第 4 行到 B2 那么多行）。')
    ws['A2'] = '挤过来的行数：'
    ws['B2'] = f'={JE}!${J_CNT}${R1}'
    ws['C2'] = '其中结转块以前：'
    ws['D2'] = f'={JE}!${J_CNT}${p1008.CLOSE0 - 1}'
    for col, t in [('A', '记账分录\n第几行')] + [(c, t) for c, _, t, _ in CMP_COLS]:
        put(ws, f'{col}3', t, F_NOTE, FILL_AUTO, align=ACW)
        ws.column_dimensions[col].width = 11
    cnt = f'{JE}!${J_CNT}${R0}:${J_CNT}${R1}'
    for i in range(N):
        r = CMP_R0 + i
        ws[f'A{r}'] = f'=IF(ROW()-{CMP_R0 - 1}>$B$2,"",IFERROR(MATCH(ROW()-{CMP_R0},{cnt},1),0)+1)'
        for c, src, _, _ in CMP_COLS:
            f = f'INDEX({JE}!${src}${R0}:${src}${R1},$A{r})'
            if src == J_CP:
                f += '&""'
            ws[f'{c}{r}'] = f'=IF($A{r}="","",{f})'
    C1 = CMP_R0 + N - 1
    # 精简…＝挤过来的全部行（科目余额表用：结转块的分录要进余额）；前段…＝只到结转块以前（往来、分类键、损益键在结转块里都是空的，
    # 往来余额表、经营报表走势、_结转损益 用这个，少扫 1,958 行；_结转损益 只能用这个范围也是为了不成循环）
    for c, _, _, nm in CMP_COLS:
        for name, cell in ((nm, '$B$2'), (nm.replace('精简', '前段'), '$D$2')):
            ref = f"{q(SH_CMP)}!${c}${CMP_R0}:INDEX({q(SH_CMP)}!${c}${CMP_R0}:${c}${C1},MAX(1,{q(SH_CMP)}!{cell}))"
            wb.defined_names[name] = DefinedName(name, attr_text=ref)
    ws.freeze_panes = 'A4'
    log.append(f'分录精简：新加隐藏表 {SH_CMP}（{N} 行 × {len(CMP_COLS) + 1} 列，有凭证的分录按顺序挤在前面），记账分录加隐藏 {J_CNT} 列（挤过来的累计行数）')

    # ── 改成只扫精简表 ──
    n = 0
    tb = wb[SH_TB]
    names3 = [('分录余额键', '精简余额键'), ('分录借方', '精简借方'), ('分录贷方', '精简贷方')]
    for r in range(TB_R0, TB_R1 + 1):
        _swap(tb, f'X{r}', [(f'ISNUMBER(MATCH($Q{r},分录科目行,0))',
                             f'COUNTIFS(精简余额键,">"&$Q{r}*{TB_KEYM},精简余额键,"<"&($Q{r}+1)*{TB_KEYM})>0')])
        for c in ('Z', 'AA', 'AB', 'AC', 'AD'):
            w = _swap(tb, f'{c}{r}', names3)
            assert '分录' not in w, w
        n += 6
    _swap(tb, 'X5', [('COUNTIF(分录余额键,"#*")', 'COUNTIF(精简余额键,"#*")')])
    log.append(f'科目余额表：X（有没有分录）、Z～AD（本月 / 本年 / 年初净额）、X5 只扫精简表（{n + 1} 格；'
               f'H～K、AE 里「科目表里找不到的编码」那一段照旧扫全表，只在 X5＞0 时才算）')

    ar = wb[SH_ARAP]
    names4 = [('分录借方', '前段借方'), ('分录贷方', '前段贷方'), ('分录往来', '前段往来'), ('分录分类键', '前段分类键')]
    m = 0
    for r in range(5, 5 + CP_R1 - CP_R0 + 1):
        _swap(ar, f'Q{r}', [(f'COUNTIF(分录往来,$A{r})', f'COUNTIF(前段往来,$A{r})')])
        m += 1
        for c in 'CDEFGHI':
            w = _swap(ar, f'{c}{r}', names4)
            assert '分录' not in w, w
            m += 1
    rt = 5 + CP_R1 - CP_R0 + 1
    assert ar[f'A{rt}'].value == '没写往来单位的', ar[f'A{rt}'].value
    for c in 'CDEFGHI':
        w = _swap(ar, f'{c}{rt}', names4)
        assert '分录' not in w, w
        m += 1
    log.append(f'往来余额表：Q（有没有分录）、C～I 各科目余额、「没写往来单位的」只扫精简表（{m} 格）')

    hp = wb[SH_CLOSE]
    rr_old = f'{JE}!$R${R0}:$R${p1008.JE1}'
    ss_old = f'{JE}!$S${R0}:$S${p1008.JE1}'
    k = 0
    for r in range(5, 5 + p1008.NPL):
        w = _swap(hp, f'E{r}', [(rr_old, '前段损益键')])
        assert '记账分录' not in w, w
        k += 1
        for i in range(p1008.NM):
            c = p1008.rawcol(i + 1)
            w = _swap(hp, f'{c}{r}', [(rr_old, '前段损益键'), (ss_old, '前段损益额')])
            assert '记账分录' not in w, w
            k += 1
    log.append(f'_结转损益：E（分录条数）、BU～CZ（各月发生额）只扫精简表（{k} 格；结转损益块本来就没有损益键）')

    ops = wb[SH_OPS]
    o = 0
    for row in ops.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith('=') and '分录分类键' in v:
                rest = v.replace('分录分类键', '').replace('分录借方', '').replace('分录贷方', '')
                if '分录' in rest:
                    continue                         # 还用到别的分录列（科目、日期、来源）的不动
                _swap(ops, c.coordinate, [('分录分类键', '前段分类键'), ('分录借方', '前段借方'), ('分录贷方', '前段贷方')])
                o += 1
    assert o >= 12 * 5, o
    log.append(f'经营报表：12 个月走势（{o} 格）只扫精简表')
    return log


def step_vx_rank(wb, ctx):
    """_凭证索引 I 列（当月第几号）：原来每张凭证 COUNTIFS(月份＝本月, 排序键≤本张) 两个条件扫 6,895 行；
    排序键＝日期×10^7＋凭证ID，同一个月的都夹在「本月 1 日×10^7」和本张之间，所以＝COUNTIF(排序键≤本张)−本月 1 日以前的张数，
    后者每个月只算一次（_分录精简 K、L 列，名称「凭证月前张数」）。"""
    import re
    log = []
    ws = wb[SH_CMP]
    IXN = q(SH_VIDX)
    E1 = ctx['vx_end']
    KY = f'{IXN}!$E$4:$E${E1}'
    NM = p1008.NM
    put(ws, 'K3', '月序', F_NOTE, FILL_AUTO, align=ACW)
    put(ws, 'L3', '本月 1 日以前\n的凭证张数', F_NOTE, FILL_AUTO, align=ACW)
    ws.column_dimensions['L'].width = 13
    for i in range(NM):
        r = CMP_R0 + i
        ws[f'K{r}'] = i + 1
        ws[f'L{r}'] = f'=COUNTIF({KY},"<"&DATE(YEAR(起始月份),MONTH(起始月份)+{i},1)*10^7)'
    from openpyxl.workbook.defined_name import DefinedName
    wb.defined_names['凭证月前张数'] = DefinedName('凭证月前张数', attr_text=f"{q(SH_CMP)}!$L${CMP_R0}:$L${CMP_R0 + NM - 1}")
    vx = wb[SH_VIDX]
    pat = re.compile(r"^=IF\(\$C(\d+)=0,\"\",COUNTIFS\('?_凭证索引'?!\$H\$4:\$H\$(\d+),\$H(\d+),'?_凭证索引'?!\$E\$4:\$E\$(\d+),\"<=\"&\$E(\d+)\)\)$")
    n = 0
    for r in range(4, E1 + 1):
        v = vx[f'I{r}'].value
        m = pat.match(v)
        assert m and all(int(x) == r for x in (m.group(1), m.group(3), m.group(5))) and int(m.group(2)) == E1 == int(m.group(4)), (r, v)
        H = f'$H{r}'
        mn = f'((INT({H}/100)-YEAR(起始月份))*12+MOD({H},100)-MONTH(起始月份)+1)'
        base = (f'IF(OR({mn}<1,{mn}>{NM}),COUNTIF({KY},"<"&DATE(INT({H}/100),MOD({H},100),1)*10^7),'
                f'INDEX(凭证月前张数,{mn}))')
        vx[f'I{r}'] = f'=IF($C{r}=0,"",COUNTIF({KY},"<="&$E{r})-{base})'
        n += 1
    log.append(f'_凭证索引 I 列（当月第几号，{n} 格）：COUNTIFS 两个条件 → COUNTIF 一个条件减「本月 1 日以前的张数」（每月只算一次）')
    return log


ROWS_CHECK = [(SH_CASH, CASH_R0, CASH_R1), (SH_BUY, MOD_R0, BUY_R1), (SH_ISS, MOD_R0, ISS_R1), (SH_SALE, MOD_R0, SALE_R1),
              (SH_PAY, MOD_R0, PAY_R1), (SH_MAN, MOD_R0, MAN_R1)]


def step_rows_check(wb):
    """数据校验加第 41 项：录入表有没有被插入 / 删除整行。
    记账分录、凭证索引现在直接引用流水 / 业务表的第几行（原来 INDEX(整列,第几笔)），删了整行，后面的凭证号会对不上——
    规矩本来就是「不要插行删行」，这里加一道检查：表的行数（ROWS）一变就报 ✗。"""
    from copy import copy
    ck = wb[SH_CHK]
    assert ck['A44'].value == 40 and ck['A45'].value is None, (ck['A44'].value, ck['A45'].value)
    bad = '&'.join(f'IF(ROWS({q(sh)}!$A${r0}:$A${r1})<>{r1 - r0 + 1},"【{sh}】","")' for sh, r0, r1 in ROWS_CHECK)
    vals = (41, '录入表：没有插入 / 删除整行',
            f'=IF({bad}="","√ 没有","✗ "&{bad}&"插入或删除了整行")',
            '记账分录、凭证直接按位置取各录入表的第几行：插入或删除整行会让后面的凭证号、凭证汇总对不上。按 Ctrl+Z 撤销；'
            '要删一笔，选中那一行的内容按 Delete 清空，不要删整行；新的往下面空行填')
    ck.row_dimensions[45].height = ck.row_dimensions[44].height
    for col, v in zip('ABCD', vals):
        src, dst = ck[f'{col}44'], ck[f'{col}45']
        dst.font, dst.fill, dst.border = copy(src.font), copy(src.fill), copy(src.border)
        dst.alignment, dst.number_format = copy(src.alignment), src.number_format
        dst.value = v
    old = ck['C3'].value
    assert old.count('C5:C44') == 3, old
    ck['C3'].value = old.replace('C5:C44', 'C5:C45')
    assert ck.auto_filter.ref == 'A4:D44', ck.auto_filter.ref
    ck.auto_filter.ref = 'A4:D45'
    from openpyxl.formatting.formatting import ConditionalFormattingList
    cf_new = ConditionalFormattingList()
    for rg in ck.conditional_formatting:
        ref = str(rg.sqref).replace('C5:C44', 'C5:C45')
        for rule in rg.rules:
            cf_new.add(ref, rule)
    ck.conditional_formatting = cf_new
    hm = wb[SH_HOME]
    n = 0
    for row in hm.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith('=') and '数据校验!$C$5:$C$44' in c.value:
                c.value = c.value.replace('数据校验!$C$5:$C$44', '数据校验!$C$5:$C$45')
                n += 1
    assert n >= 1, '首页没找到数据校验汇总格'
    return [f'数据校验加第 41 项「录入表没有插入 / 删除整行」（首页汇总 {n} 格范围跟着扩到第 45 行）']
