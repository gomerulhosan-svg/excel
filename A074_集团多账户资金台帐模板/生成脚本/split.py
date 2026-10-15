# -*- coding: utf-8 -*-
"""a74split 拆本：在 build 的内存里，从一本版的 wb（make.build）出「查询本」「录入本」两本。做法照 A065、A075（A075 split.py 只读参考）：
外部链接在 openpyxl 最后一次存盘时建（读回成品再存会丢外部链接关系、丢合并区里的公式），之后只做 zip 级处理（share、inject、刷外部缓存）。

    undo = split.make_query(wb, ctx)   →  make.check_book(wb, 'query')  →  wb.save(查询本)  →  undo()
    split.make_entry(wb, ctx)          →  make.check_book(wb, 'entry')  →  wb.save(录入本)          （录入本是拆掉的，最后做）

查询本（make_query，就地改，返回 undo）：一本版原样（报表、_流、段表、块表、分段名称 台k_字段 一个字不改），只改：
  · 【资金台帐】【基础资料】变成同名同位置的隐藏取数表：录入格换成取数公式 mirror_formula()，显示列（台帐 H、L..O，基础资料 序号、核对）照一本版算；
    台帐取 A..G、I..K、P 第 HDR～FETCH_LAST 行（表头那一行也取，查「录入本表头挪了」）；基础资料取①的 FETCH_PARAMS、②～⑥各块的填写列（表头行～块尾）。
    去掉这两张的有效性、条件格式；
  · _参：P_录入模式 恒为 0；P_锚坏 改成「锚列/段表坏了，或者取过来的录入本表头不是原来的字」（录入本表头上面插删了行、插删了列、基础资料插删了行）；
  · 首页 K3 改成「数据来自录入本：流水 N 笔，最后一笔 yyyy/m/d」；横幅、核对表「台帐分段」里锚坏那句话换成「录入本的表头挪了」；
    指向隐藏取数表的超链接去掉（导航格写「（在录入本）」）；【使用说明】换成查询本版；
  · 外部链接（openpyxl ExternalLink，目标＝相对文件名 WB1_FILE）。
录入本（make_entry，就地拆）：只留 ENTRY_SHEETS＋ENTRY_HIDDEN；资金台帐显示列、辅助列、第 1/3 行换成 s_entry 的轻公式，锚列 Q 清掉（列位不挪）；
  基础资料①「录入模式」那一行改成一句说明；_参 只留轻的键（再加 挪了/插行，可录 改成录入本口径）；_表 去掉报表才用的几组（分类候选、两级、月份）；
  新建 _块、首页（轻）、使用说明（录入本版）；名称只留用得上的。"""
import re

from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidationList

import layout as Lt
from layout import (
    SH_HOME, SH_LEDGER, SH_BASE, SH_HELP, SH_CHK, H_PAR, H_TAB, H_FLOW, HDR, R0, N_LEDGER, LEDGER_COLS, LEDGER_SHOW, LEDGER_HEAD,
    LEDGER_LAST, LEDGER_LAST_ROW, LEDGER_KPI, LEDGER_ALERT, LEDGER_ANCHOR, BASE_BLOCKS, BASE_COLS, BASE_HEAD, BASE_LAST, PAR,
    PAR_KEYS, PAR_ROW, BASE_PAR_VAL, HOME_SEL, WB1_FILE, FETCH_LAST, FETCH_PARAMS, EXT_PREFIX, SH_BLK_E, ENTRY_SHEETS,
    ENTRY_HIDDEN, ENTRY_HELP, ENTRY_HELP_NAME, ENTRY_KPI, ENTRY_KPI_NOTE, _TAB, q, inref)
from common import F_HELP, F_AUTO, F_KPI_L, F_KPI_V, F_NOTE, FILL_AUTO, FILL_KPI, FILL_NONE, BD, AC, ACW, ALW, put, kpi
import s_entry
import s_input
import s_home

# ═════════════════════════════ 取数公式 ═════════════════════════════
def _xref(sh, a):
    """'[1]表名'!A5（外部引用一律加引号）"""
    return "'[1]" + sh.replace("'", "''") + "'!" + a


def mirror_formula(sh, col, r, name=None, r0=None):
    """取数公式。MIRROR_MODE='name'（给了 name/r0）＝INDEX(名称,ROW()-常数)；否则每格直接链接（默认）。
       IF(x="","",x)：录入本空着的格子给 ""（给不出真的空格；一本版里会把 ""/空格 当成不一样的写法已经改掉，见 layout.cnta）"""
    x = f'INDEX({name},ROW()-{r0 - 1})' if name else _xref(sh, f'{col}{r}')
    f = f'IF({x}="","",{x})'
    return f'=IFERROR({f},"")' if Lt.MIRROR_IFERROR else '=' + f


def fetch_spec():
    """[(表, 列, 起行, 止行, 名称标签)]：查询本里要换成取数公式的格子（列 × 连续行）"""
    out = [(SH_LEDGER, c, HDR, FETCH_LAST, f'台帐_{c}') for c in LEDGER_COLS.values()]
    rows = sorted(PAR_ROW[k] for k in FETCH_PARAMS)
    assert rows == list(range(rows[0], rows[-1] + 1))
    out.append((SH_BASE, BASE_PAR_VAL, rows[0], rows[-1], f'基_参_{BASE_PAR_VAL}'))
    for k, _t, _tr, hdr, _cap in BASE_BLOCKS:
        for f, c in BASE_COLS[k].items():
            if f in ('序号', '核对'):
                continue
            out.append((SH_BASE, c, hdr, BASE_LAST[k], f'基_{k}_{c}'))
    return out


def add_external_link(wb, target, sheet_names):
    """查询本 → 录入本的外部链接（照 A065/A075 add_external_link，目标写相对文件名）。返回链接对象（undo 用）"""
    from openpyxl.workbook.external_link.external import (ExternalLink, ExternalBook, ExternalSheetNames,
                                                           ExternalSheetDataSet, ExternalSheetData)
    from openpyxl.packaging.relationship import Relationship
    eb = ExternalBook(sheetNames=ExternalSheetNames(sheetName=list(sheet_names)),
                      sheetDataSet=ExternalSheetDataSet(sheetData=[ExternalSheetData(sheetId=i) for i in range(len(sheet_names))]),
                      id='rId1')
    el = ExternalLink(externalBook=eb)
    el.file_link = Relationship(type='externalLinkPath', Target=target, TargetMode='External', Id='rId1')
    assert not wb._external_links, '一本版里已经有外部链接了'
    wb._external_links.append(el)
    return el


# ═════════════════════════════ 查询本 ═════════════════════════════
Q_ANCHOR = ('※ 录入本的表头挪了（资金台帐表头上面、或基础资料里插删了行，或者插删了列）：这本按位置读，读错了。'
            '回录入本撤销（Ctrl+Z）；基础资料别在中间插行，在各块下面空行加')
Q_ANCHOR_CHK = ('"录入本的表头挪了：资金台帐第 4 行、基础资料各块表头不在原来那一行（录入本表头上面、基础资料中间插删了行），或者插删了列。'
                '这本按位置读，所有数都读错了。回录入本撤销（Ctrl+Z）再存盘；基础资料加东西在各块下面的空行加"')
_CHK_TEXT = [   # 核对表「台帐分段」：一本版说锚列/表头被删的话 → 查询本的说法
    ('"台帐的表头（第 4 行）或右边的 Q 列（勿删）被删了：所有数都没算。马上撤销（Ctrl+Z）；已经存盘了找我们（重出一本模板，把数导过去）"',
     Q_ANCHOR_CHK),
    ('"✗ 表头或 Q 列被删"', '"✗ 录入本表头挪了"'),
    ('"（算不出来：多半是 Q 列或表头被删了，马上撤销）"', '"（算不出来：多半是录入本表头挪了）"'),
]


def query_anchor_formula():
    """查询本 P_锚坏：段表坏了，或者取过来的录入本表头（台帐第 4 行的录入列、基础资料各块表头的第一列和最后一个填写列）不是原来的字"""
    L, B = q(SH_LEDGER), q(SH_BASE)
    first = lambda t: t.split('\n')[0]
    cond = []
    for f, col in LEDGER_COLS.items():
        t = first(LEDGER_HEAD[f])
        cond.append(f'LEFT({L}!${col}${HDR}&"",{len(t)})<>"{t}"')
    for k, _t, _tr, hdr, _cap in BASE_BLOCKS:
        cols = [(f, c) for f, c in BASE_COLS[k].items() if f not in ('序号', '核对')]
        for f, c in (cols[0], cols[-1]):
            t = first(s_input._HEAD_FIX.get((k, f), BASE_HEAD[k][f]))
            cond.append(f'LEFT({B}!${c}${hdr}&"",{len(t)})<>"{t}"')
    return f'=IFERROR(IF(OR(ISERROR(INDEX(段_首,1)),COUNT(段_末)=0,{",".join(cond)}),1,0),1)'


def home_k3_formula():
    d = 'MAX(流_日期)'
    return (f'="数据来自录入本：流水 "&TEXT(P_已用行数,"#,##0")&" 笔，最后一笔 "&IF({d}>0,{s_home.dtxt(d)},"（没有）")'
            f'&"（跟录入本首页对一下）"')


class Undo:
    def __init__(self):
        self.ops = []

    def set(self, ws, a, v):
        c = ws[a]
        self.ops.append(('v', c, c._value, c.data_type))
        c.value = v

    def hyper(self, c, v):
        self.ops.append(('h', c, c._hyperlink))
        c._hyperlink = v

    def attr(self, obj, name, v):
        self.ops.append(('a', obj, name, getattr(obj, name)))
        setattr(obj, name, v)

    def call(self, fn):
        self.ops.append(('f', fn))

    def __call__(self):
        for op in reversed(self.ops):
            if op[0] == 'v':
                _k, c, v, dt = op
                c._value, c.data_type = v, dt
            elif op[0] == 'h':
                op[1]._hyperlink = op[2]
            elif op[0] == 'a':
                setattr(op[1], op[2], op[3])
            else:
                op[1]()
        self.ops = []


def _repl_in_formulas(u, ws, pairs, need=True):
    n = 0
    for c in list(ws._cells.values()):                  # 只看有的格子（iter_rows 会把空位都建出来）
        v = c.value
        if isinstance(v, str) and v.startswith('='):
            w = v
            for a, b in pairs:
                w = w.replace(a, b)
            if w != v:
                u.set(ws, c.coordinate, w)
                n += 1
    assert n or not need, f'{ws.title}：要换的字一处也没找到（{pairs[0][0][:30]}…）'
    return n


def make_query(wb, ctx, log=print):
    mode = Lt.MIRROR_MODE
    assert mode in ('cell', 'name'), mode
    u = Undo()
    n_cell = {}
    new_names = []
    for sh, col, r0, r1, tag in fetch_spec():
        ws = wb[sh]
        nm = None
        if mode == 'name':
            nm = f'{EXT_PREFIX}{tag}'
            assert nm not in wb.defined_names
            # 名称里的外部引用：表名不用引号时不加（跟 Excel 自己写的一样；LibreOffice 连着更新时加了引号的外部名称取不到数）
            ext = (f'[1]{sh}!' if not q(sh).startswith("'") else _xref(sh, '')) + f'${col}${r0}:${col}${r1}'
            wb.defined_names[nm] = DefinedName(nm, attr_text=ext, hidden=True)
            new_names.append(nm)
        for r in range(r0, r1 + 1):
            a = f'{col}{r}'
            c = ws[a]
            assert not isinstance(c, MergedCell), f'{sh}!{a} 在合并区里'
            assert not (isinstance(c.value, str) and c.value.startswith('=')), f'{sh}!{a} 是公式，不该取数'
            u.set(ws, a, mirror_formula(sh, col, r, nm, r0))
            n_cell[sh] = n_cell.get(sh, 0) + 1
    for sh in (SH_LEDGER, SH_BASE):
        ws = wb[sh]
        u.attr(ws, 'sheet_state', 'hidden')
        u.attr(ws, 'data_validations', DataValidationList())
        u.attr(ws, 'conditional_formatting', ConditionalFormattingList())
    # _参
    par = wb[H_PAR]
    u.set(par, PAR['录入模式'], 0)
    u.set(par, PAR['锚坏'], query_anchor_formula())
    u.set(par, f'C{PAR["录入模式"][1:]}', '查询本：恒为 0（两本版不用录入模式）')
    u.set(par, f'C{PAR["锚坏"][1:]}', '查询本：1＝段表坏了，或者取过来的录入本表头不是原来的字（录入本插删了表头上面的行、基础资料中间插删了行、插删了列）')
    # 首页：K3、横幅
    home = wb[SH_HOME]
    u.set(home, 'K3', home_k3_formula())
    _repl_in_formulas(u, home, [(s_home.B_ANCHOR, Q_ANCHOR)])
    _repl_in_formulas(u, wb[SH_CHK], _CHK_TEXT)
    # 指向隐藏取数表的超链接去掉；首页导航格写「（在录入本）」
    hid = {SH_LEDGER, SH_BASE}
    n_link = 0
    for ws in wb.worksheets:
        if ws.title in hid or ws.sheet_state != 'visible':
            continue
        for c in list(ws._cells.values()):
            h = getattr(c, '_hyperlink', None)
            if h is None or not h.location:
                continue
            m = re.match(r"^'?((?:[^']|'')+?)'?!", h.location)
            if m and m.group(1).replace("''", "'") in hid:
                u.hyper(c, None)
                n_link += 1
                if ws.title == SH_HOME and c.value in hid:
                    u.set(ws, c.coordinate, f'{c.value}（在录入本）')
    # 使用说明 → 查询本版（新表换进同一个位置；undo 换回来）
    old = wb[SH_HELP]
    idx = wb._sheets.index(old)
    new = wb.create_sheet(SH_HELP + '_q')
    wb._sheets.remove(new)
    wb._sheets[idx] = new
    new.title = SH_HELP
    new.sheet_properties.tabColor = old.sheet_properties.tabColor
    s_entry.build_help_query(new, ctx)

    def back_help():
        wb._sheets[idx] = old
    u.call(back_help)
    # 外部链接（表名清单＝录入本的表；refresh_link_cache 按录入本实际的重写）
    link = add_external_link(wb, WB1_FILE, ENTRY_SHEETS + ENTRY_HIDDEN)

    def back_link():
        wb._external_links.remove(link)
        for nm in new_names:
            del wb.defined_names[nm]
    u.call(back_link)
    log(f'   · 查询本：取数公式 {sum(n_cell.values()):,} 格（写法 {mode}' + (f'，外部名称 {len(new_names)} 个' if new_names else '') + '）：'
        + '、'.join(f'{k} {v:,}' for k, v in n_cell.items()) + f'；去掉指向取数表的链接 {n_link} 个')
    return u


# ═════════════════════════════ 录入本 ═════════════════════════════
E_PAR_KEEP = ('集团名称', '开账日', '开账年', '日期上限', '已用行数', '期初合计', '待定号', '内部号', '账户数', '公司数', '收费项数',
              '收付人数', '规则数', '项规则数', '人规则数', 'Q规则数', '可录')
E_TAB_DROP = ('分类候选', '两级', '月份')
ENTRY_TIP = (
    '★ 录入本：只在这里录，报表、子表、核对表在《' + Lt.WB2_FILE[:-5] + '》（从这本取数）。一笔一行，日期、账户必填（淡黄）；收入、支出只填一个'
    '（网银的「−5000」照样粘进收入栏，会自动当支出）。内部转账记两行，两行的「转账的另一户」都填对方，收费项空着会自动当内部转账。'
    '收费项空着会按摘要认（N 列琥珀色斜体）；一笔钱拆成几行的一定要自己选。粘网银用「选择性粘贴→数值」，只粘 A～G、I～K、P 列。'
    '灰色几列（H、L～O）自动算，别动；L 列 ※ 必须改，△ 看一眼（这里只查这一行看得出来的毛病，配对、跳期、重复粘贴看查询本的核对表）。'
    '别插行：漏记的录在最下面就行，日期写对，查询本的报表按日期算；调顺序用日期列的筛选箭头排序。不要的行点行号删整行；'
    '只清内容就只选 A～G、I～K、P 按 Delete（别整行按 Delete，会把右边的自动列也清掉）。别在表头上面插删行、别插删列（查询本按位置读）。'
    '录完存盘；两本别同时开着录。')
BASE_TIP_E = (
    '★ 名单和参数都在这里：① 参数 ② 公司 ③ 账户 ④ 收费项 ⑤ 收付人 ⑥ 摘要关键词规则（点第 3 行跳过去）。淡黄格可以改，灰色是自动的（序号、核对）。'
    '名单往各块下面的空行接着加，别在中间插行、删行（查询本按位置读，插删了会读错）；③可以排序（子表靠「子表名」认账户，那一列别改）。'
    '账户、收费项、收付人改名时，把旧名字加进「别名」（用「、」隔开），台帐里旧的写法照样认。'
    '不用的账户把状态改成「销户」、填停用日，别删行。年中新开的账户期初余额填 0，开户那天的余额在台帐录一笔「账户转入」。核对列 ※ 必须改，△ 看一眼。')


def _names_in(f, names):
    body = ''.join(f.split('"')[0::2])
    return {t for t in re.findall(r"[A-Za-z_\\一-鿿][\w.一-鿿]*", body) if t in names}


def make_entry(wb, ctx, log=print):
    keep = set(ENTRY_SHEETS) | {H_PAR, H_TAB}
    for ws in list(wb.worksheets):
        if ws.title not in keep:
            del wb[ws.title]
    for sh in (SH_HOME, SH_HELP):                     # 一本版的首页、使用说明删掉，下面重建
        del wb[sh]
    led, base, par, tab = wb[SH_LEDGER], wb[SH_BASE], wb[H_PAR], wb[H_TAB]
    # —— 资金台帐：显示列、辅助列、第 1/2/3 行、锚列 ——
    for r in range(R0, LEDGER_LAST_ROW + 1):
        sf = s_entry.show_formulas(r)
        for k, col in LEDGER_SHOW.items():
            led[f'{col}{r}'].value = sf[k]
        hf = s_entry.help_formulas(r)
        for k, col in ENTRY_HELP.items():
            c = led[f'{col}{r}']
            c.value = hf[k]
            c.font = F_HELP
    for k, col in ENTRY_HELP.items():
        c = led[f'{col}{HDR}']
        c.value = f'{k}（自动，勿动）'
        c.font = F_HELP
        led.column_dimensions[col].hidden = True
    # 锚列 Q：录入本不要了（空着，列位不挪）：表头「勿删」、列样式、有效性去掉
    led[f'{LEDGER_ANCHOR}{HDR}'].value = None
    led[f'{LEDGER_ANCHOR}{HDR}'].style = 'Normal'
    cd = led.column_dimensions[LEDGER_ANCHOR]
    cd.width = 3
    cd.font = Font(name=Lt.FONT, sz=10)
    from openpyxl.styles import PatternFill, Border
    cd.fill = PatternFill(fill_type=None)
    cd.border = Border()
    dvs = DataValidationList()
    for dv in led.data_validations.dataValidation:
        if str(dv.sqref).startswith(f'{LEDGER_ANCHOR}{HDR}'):
            continue
        if dv.prompt and '往下拖一格' in dv.prompt:
            dv.prompt = '灰色是自动算的，别在这里打字、粘贴。别插行：漏记的录在最下面就行'
        dvs.append(dv)
    led.data_validations = dvs
    led[LEDGER_ALERT].value = s_entry.alert_formula()
    led['A2'].value = ENTRY_TIP
    kf = s_entry.kpi_formulas()
    for (a, b, _l) in LEDGER_KPI:                        # 一本版的汇总条格子先清掉
        led[a].value = led[b].value = None
    for a, b, k in ENTRY_KPI:
        kpi(led, a, b, k, kf[k], fmt=s_entry.KPI_FMT[k])
        if k in ('※ 行数', '△ 行数', '还能录'):
            led[b].font = Font(name=Lt.FONT, sz=11, bold=True, color='FFC00000')
    nc = ENTRY_KPI_NOTE
    led.merge_cells(f'{nc}:{LEDGER_LAST}3')
    put(led, nc, s_entry.KPI_NOTE, F_NOTE, FILL_KPI, align=ALW)
    led.conditional_formatting = _cf_without(led.conditional_formatting, {f'{c}3' for _a, c, _l in LEDGER_KPI} |
                                             {a for a, _c, _l in LEDGER_KPI})
    # —— 基础资料：①「录入模式」那一行改成一句说明；①「管理年度」的「现在按」不引用查询本才有的 P_管理年度；说明 ——
    r = PAR_ROW['录入模式']
    base[f'A{r}'].value = '（两本版不用）'
    base[f'B{r}'].value = None
    base[f'B{r}'].fill = FILL_AUTO
    base[f'C{r}'].value = '原来的「录入模式」：两本版不用了——录的时候只开这本录入本，本来就不卡；报表、子表、核对表在查询本里算'
    base[f'L{r}'].value = None
    rm = PAR_ROW['管理年度']
    b = f'${BASE_PAR_VAL}${rm}'
    base[f'L{rm}'].value = (f'=IF({Lt.raw(b)},IF(AND(ISNUMBER({b}),{b}>=2000,{b}<=2100),"查询本按 "&{b}&" 年看","※ 看不懂：填 4 位年份"),'
                            f'"空着＝查询本按最后一笔流水的年份看")')
    base['A2'].value = BASE_TIP_E
    dvs = DataValidationList()
    for dv in base.data_validations.dataValidation:
        if str(dv.sqref) == f'{BASE_PAR_VAL}{r}':
            continue
        dvs.append(dv)
    base.data_validations = dvs
    # —— _参：只留轻的键；可录改成录入本口径；加 挪了、插行 ——
    gone_par = []
    for k in PAR_KEYS:
        if k in E_PAR_KEEP:
            continue
        a = PAR[k]
        par[a].value = None
        par[f'C{a[1:]}'].value = '（录入本不用）'
        gone_par.append(f'P_{k}')
    par[PAR['可录']].value = f'=MAX(0,{N_LEDGER}-P_已用行数)'
    par[f'C{PAR["可录"][1:]}'].value = f'录入本：还能录约几行（台帐 {N_LEDGER:,} 行）'
    rr = len(PAR_KEYS) + 3
    new_names = {}
    for k, f, note in s_entry.entry_par_rows():
        par[f'A{rr}'], par[f'B{rr}'], par[f'C{rr}'] = k, f, note
        new_names[f'P_{k}'] = f'{H_PAR}!$B${rr}'
        rr += 1
    # —— _表：报表才用的几组清掉 ——
    drop = {CI(col) for g in E_TAB_DROP for col in _TAB[g].values()}
    for (r_, c_), c in list(tab._cells.items()):
        if c_ in drop:
            del tab._cells[(r_, c_)]
    # —— _块、首页、使用说明 ——
    blk = wb.create_sheet(SH_BLK_E)
    s_entry.build_blk(blk)
    home = wb.create_sheet(SH_HOME)
    home.sheet_properties.tabColor = Lt.C_HOME[2:]
    s_entry.build_home(home, ctx)
    hp = wb.create_sheet(SH_HELP)
    hp.sheet_properties.tabColor = Lt.C_HOME[2:]
    s_entry.build_help_entry(hp, ctx)
    # —— 名称：删掉指向删掉的表 / 录入本不用的 P_；加录入本自己的；再只留用得上的 ——
    sheets = set(wb.sheetnames)
    for nm in list(wb.defined_names):
        t = wb.defined_names[nm].attr_text or ''
        refs = re.findall(r"(?:'((?:[^']|'')+)'|([^\W\d][\w.]*))!", t)
        if any((a or b).replace("''", "'") not in sheets for a, b in refs) or nm in gone_par:
            del wb.defined_names[nm]
    for nm, col in ENTRY_HELP_NAME.items():
        new_names[col] = inref(SH_LEDGER, ENTRY_HELP[nm])
    new_names.update(s_entry.blk_names())
    for nm, ref in new_names.items():
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)
    names = {n: dn.attr_text or '' for n, dn in wb.defined_names.items()}
    used, stack = set(), []
    for ws in wb.worksheets:
        fs = [c.value for c in ws._cells.values() if isinstance(c.value, str) and c.value.startswith('=')]
        fs += ['=' + (dv.formula1 or '') for dv in ws.data_validations.dataValidation]
        fs += ['=' + f for cf in ws.conditional_formatting for rule in cf.rules for f in (rule.formula or [])]
        for f in fs:
            stack += list(_names_in(f, names))
    while stack:
        t = stack.pop()
        if t in used:
            continue
        used.add(t)
        stack += list(_names_in('=' + names[t], names))
    n0 = len(names)
    for nm in names:
        if nm not in used:
            del wb.defined_names[nm]
    # —— 表顺序、隐藏、活动表 ——
    order = ENTRY_SHEETS + ENTRY_HIDDEN
    assert set(order) == set(wb.sheetnames), (order, wb.sheetnames)
    wb._sheets = [wb[s] for s in order]
    for s in ENTRY_HIDDEN:
        wb[s].sheet_state = 'hidden'
    for s in ENTRY_SHEETS:
        wb[s].sheet_state = 'visible'
    wb.active = ENTRY_SHEETS.index(SH_LEDGER)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = ws.title == SH_LEDGER
    for ws in wb.worksheets:
        ws.__dict__.pop('_a074_f', None)
    log(f'   · 录入本：{len(wb.sheetnames)} 张表（{"、".join(wb.sheetnames)}）；名称 {n0:,} → {len(used):,} 个')


def _cf_without(cfl, cells):
    """条件格式里去掉只管这几格的规则（汇总条换了内容）"""
    out = ConditionalFormattingList()
    for cf in cfl:
        rng = str(cf.sqref)
        if rng in cells:
            continue
        for rule in cf.rules:
            out.add(rng, rule)
    return out
