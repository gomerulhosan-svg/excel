# -*- coding: utf-8 -*-
"""A074 集团多账户资金台帐 · 隐藏计算层 v3.1（设计.md §2；提速改动见 §15.1、§15.3）：
  _参  全书参数（B 列，名称 P_xx；v3.1 加 台帐末/台帐容量/漏行/Q规则数 和分段报警 P_错位/笔数差/没算/锚坏/可录）
  _表  基础资料镜像：取基础资料一律用名称 基_块_字段（layout.bref），不直接写「基础资料!」（WPS 存盘会拆）；
       压紧倒序的规则表（项规则、人规则、能在收付人里找的项规则 规Q）＋下拉压紧＋收支分类候选键＋公司→账户两级候选键
       ＋月份下拉清单（change 10/14：MONTH_LIST 第 2..13 行＝管理年度 1～12 月的文字「2026年1月」，名称 月份_清单只给下拉用）
  _流  资金台帐分段取数（滚动接力）：台帐每 SEG_ROWS 行一段，_流 每段 SEG_SLOTS 个位置；第 z 个位置在第 z+1 行。
       取台帐格只用本段读数区名称 台k_字段（layout.sref）；位置跟台帐行号的对应在段表（GY:HM）。
       块表（一块＝一段：各账户净额/块前累计/最晚日期，窗口用的月统计）、段表、规则推断自检格。

防环、性能（设计 §9.1、§15.1.7、§15.3）：
  原始列（取数号～月码、日期原/收值/支值/文字键/项推）只看本行、本段名称、_表、PAR_SAFE、段表那一行的链；
  正常日、余额、顺序挂在块表累计链上：直接写本块那一行的块表格（不写 INDEX(块_xx,…)），第 2 段起用定长窗口（WPS 不拆）；
  问题、提醒是短合并列；贵的 ※ 链、△ 链、重复粘贴、规则 LOOKUP 都在只看本行的 问本/提本/项推 里（改中间一行时后面几十段不重算它们）。
  反查只看段表的链（改数据永远不脏）；段表叶子（错位、有字）只给 _参 的 P_错位、P_笔数差 用。
  录入模式（P_录入模式＝1）只停 配对、期人键（FLOW_GATED）；其余照算。
"""
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from layout import (
    H_PAR, H_TAB, H_FLOW, SH_BASE, SH_HOME, SH_SPAN, PAR, PAR_KEYS, PAR_ROW, BASE_PAR_VAL,
    LIST_CO, LIST_ACC, LIST_ITEM, LIST_GROUP, LIST_PARTY, LIST_RULE, LIST_RULE_I, LIST_RULE_P, LIST_RULE_Q,
    DV_COLS, DV_CNT, CAND_ROWS, CAND_TWO, CAND_TWO_LAST, _TAB, _TAB_NEXT, HOME_SEL, SPAN_SEL, MONTH_LIST, MONTH_LIST_F,
    FLOW, FLOW_COLS, FLOW_BLK, FLOW_SEG, FLOW_PROBE, FLOW_CONST, N_FLOW, N_SEG, SEG_SLOTS, BLK, N_BLK, JUMP_LOOK,
    N_CO, N_ACC, N_ITEM, N_PARTY, N_RULE, N_RULE_P, N_RULE_Q, N_GROUP, N_YEARS, JUMP_DAYS, PAIR_DAYS, DUP_WIN,
    ATTR_NO, ATTR_UNKNOWN, ACC_STATUS, RULE_DIRS, RULE_WHERE, ITEM_PENDING, LEDGER_COLS,
    K_DAY, K_YM, K_PAIR, q, bref, lrange, esc, norm, raw, name2no, chain, dsel, win_first, win_last, win, ym_of,
    k_aypm, k_igypm, numv, cnta, tref, sref, seg_name, seg_area, read_area, seg_cell, slot_of_row, win_row, blk_row, pause, pause0, _ymd8)

R1 = 2                                  # 第 1 个位置在第 2 行
RN = N_FLOW + 1                         # 最后一个位置所在行
S = SEG_SLOTS
# 「有内容」看这 9 格（备注、原表位置不算）：取数号的门控和 段_有字 用同一个判法
GATE = ('日期', '账户', '收费项', '收付人', '摘要', '收入', '支出', '对方账户', '归属公司')
ATTR_LIST = list(ATTR_NO)               # 收入 支出 投资建设 往来 待分类 内部转账
ATTR_NUMS = [ATTR_NO[a] for a in ATTR_LIST]
INT9 = ATTR_NO['内部转账']
# _表 私有帮助列（layout 列之后，不进 NAMES）：大类首现的累计（代替往下撑的 MAX($X$1:X{r-1})，WPS 不拆）
GRP_RUN = CL(CI(_TAB_NEXT) + 1)

# 提示文字（核对表、首页按开头数数：「△ 重复粘贴*」「△ 本账户上面有日期更晚的*」「※ 比前面最晚的一笔*」「△ 比前面最晚的一笔晚*」别改开头）
T_DUP = f'△ 重复粘贴了？前后 {DUP_WIN} 行有一笔一模一样的'
T_ORDER = '△ 本账户上面有日期更晚的：按日期排一下序'
T_NEG = '△ 这一笔之后本账户余额是负数：漏录收入或顺序不对？'
# 复核 F1/data-1：冲减提示里不再出现「摘要…「」（核对表「按摘要认的」清单用「△*摘要*「*」找，会把冲减行也找进去）
T_REFUND_IN = '△ 收入类记在支出栏，按冲减收入算：是退款就在摘要或备注写个退字'
T_REFUND_OUT = '△ 支出类记在收入栏，按冲减支出算：是退款就在摘要或备注写个退字'
# 复核 F9：收付人没填的提示写上怎么改
T_NOPARTY = '△ 收付人没填，摘要也认不出：在 D 列填收付人（银行扣的费填开户行），常出现的去⑥加一条规则'


def _bx(block, field):
    """_表 第 k 行（k＝ROW()-1）对应基础资料该块第 k 个：INDEX(基_块_字段, ROW())（名称从表头行开始，第 k 个是第 k+1 格）"""
    return f'INDEX({bref(block, field)},ROW())'


def _alias(x):
    """别名规：别名空＝""；否则 "、"&norm(别名，把 ，, ;； 换成 、)&"、" """
    s = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"，","、"),",","、"),"；","、"),";","、")'
    return f'"、"&{norm(s)}&"、"'


# ═════════════════════════════ _表 ═════════════════════════════
def build_tab(ws, ctx):
    T = _TAB
    for grp, cols in T.items():
        for f, c in cols.items():
            ws[f'{c}1'] = f'{grp}·{f}'
    # —— 公司 ——
    C = LIST_CO
    for r in range(2, N_CO + 2):
        x = _bx('公司', '公司简称')
        ws[f'{C["有效"]}{r}'] = f'=IF({raw(x)},1,0)'
        ws[f'{C["名称"]}{r}'] = f'=IF({C["有效"]}{r}=1,TRIM({x}&""),"")'
        ws[f'{C["名称规"]}{r}'] = f'=IF({C["有效"]}{r}=1,{norm(x)},CHAR(1))'
    # —— 账户 ——
    A = LIST_ACC
    for r in range(2, N_ACC + 2):
        v = f'{A["有效"]}{r}'
        x = lambda f: _bx('账户', f)
        ws[v] = f'=IF({raw(x("账户名称"))},1,0)'
        ws[f'{A["名称"]}{r}'] = f'=IF({v}=1,TRIM({x("账户名称")}&""),"")'
        ws[f'{A["名称规"]}{r}'] = f'=IF({v}=1,{norm(x("账户名称"))},CHAR(1))'
        ws[f'{A["别名规"]}{r}'] = f'=IF(AND({v}=1,{raw(x("别名"))}),{_alias(x("别名"))},"")'
        ws[f'{A["子表名规"]}{r}'] = f'=IF(AND({v}=1,{raw(x("子表名"))}),{norm(x("子表名"))},CHAR(1))'
        ws[f'{A["公司号"]}{r}'] = f'=IF({v}=1,{name2no(x("公司"), "公司")},0)'
        t = f'TRIM({x("户别")}&"")'
        has = lambda w: f'ISNUMBER(SEARCH("{w}",{t}))'
        ws[f'{A["类型号"]}{r}'] = (f'=IF({v}<>1,0,IF(OR({has("对公")},{has("收单")},{has("基本户")},{has("一般户")}),1,'
                                  f'IF({has("私")},2,IF({has("现金")},3,4))))')
        ws[f'{A["期初"]}{r}'] = f'=IF({v}=1,{numv(x("期初余额"))},0)'
        st = f'TRIM({x("状态")}&"")'
        arr = '{' + ','.join(f'"{s}"' for s in ACC_STATUS) + '}'
        ws[f'{A["状态号"]}{r}'] = f'=IF({v}<>1,0,IF({st}="",1,IFERROR(MATCH({st},{arr},0),1)))'
        ws[f'{A["停用日"]}{r}'] = f'=IF({v}=1,{dsel(x("停用日"), "0")},0)'
    # —— 收费项 ——
    I = LIST_ITEM
    attr_arr = '{' + ','.join(f'"{a}"' for a in ATTR_LIST) + '}'
    fc = I['大类首现']
    ws[f'{GRP_RUN}1'] = '私·大类首现累计'
    for r in range(2, N_ITEM + 2):
        v = f'{I["有效"]}{r}'
        x = lambda f: _bx('收费项', f)
        ws[v] = f'=IF({raw(x("收费项"))},1,0)'
        ws[f'{I["名称"]}{r}'] = f'=IF({v}=1,TRIM({x("收费项")}&""),"")'
        ws[f'{I["名称规"]}{r}'] = f'=IF({v}=1,{norm(x("收费项"))},CHAR(1))'
        ws[f'{I["别名规"]}{r}'] = f'=IF(AND({v}=1,{raw(x("别名"))}),{_alias(x("别名"))},"")'
        ws[f'{I["属性号"]}{r}'] = (f'=IF({v}<>1,0,IFERROR(CHOOSE(MATCH(TRIM({x("属性")}&""),{attr_arr},0),'
                                  f'{",".join(str(n) for n in ATTR_NUMS)}),0))')
        g = f'{I["大类"]}{r}'
        ws[g] = f'=IF({v}=1,TRIM({x("大类")}&""),"")'
        # 第一次出现的大类给序号＝上面已出现的大类数＋1（累计放私有列 GRP_RUN：同表、相对引用，不写往下撑的区域）
        ws[f'{fc}{r}'] = f'=IF({g}="","",IF(MATCH({esc(g)},项_大类,0)=ROW()-1,N({GRP_RUN}{r - 1})+1,""))'
        ws[f'{GRP_RUN}{r}'] = f'=IF({fc}{r}="",N({GRP_RUN}{r - 1}),{fc}{r})'
        gn = f'{I["大类号"]}{r}'
        ws[gn] = (f'=IF({v}<>1,0,IF({g}="",99,IF(INDEX(项_大类首现,MATCH({esc(g)},项_大类,0))>{N_GROUP},99,'
                  f'INDEX(项_大类首现,MATCH({esc(g)},项_大类,0)))))')
        an = f'{I["属性号"]}{r}'
        for p in ATTR_NUMS:
            cond = f'OR({an}={p},{an}=0)' if p in (ATTR_NO['收入'], ATTR_NO['支出']) else f'{an}={p}'
            ws[f'{I[f"序键{p}"]}{r}'] = f'=IF(AND({v}=1,{cond}),{gn}*1000+ROW()-1,"")'
    for p in ATTR_NUMS:
        c = I[f'序键{p}']
        ws[f'{c}1'] = f'=COUNT({c}2:{c}{N_ITEM + 1})'
    # —— 大类压紧清单 ——
    gc = LIST_GROUP['名称']
    ws[f'{gc}1'] = '=MAX(项_大类首现)'
    for r in range(2, N_GROUP + 2):
        ws[f'{gc}{r}'] = f'=IF(ROW()-1>${gc}$1,"",INDEX(项_大类,MATCH(ROW()-1,项_大类首现,0)))'
    # —— 收付人 ——
    P = LIST_PARTY
    for r in range(2, N_PARTY + 2):
        v = f'{P["有效"]}{r}'
        x = lambda f: _bx('收付人', f)
        ws[v] = f'=IF({raw(x("收付人"))},1,0)'
        ws[f'{P["名称"]}{r}'] = f'=IF({v}=1,TRIM({x("收付人")}&""),"")'
        ws[f'{P["名称规"]}{r}'] = f'=IF({v}=1,{norm(x("收付人"))},CHAR(1))'
        ws[f'{P["别名规"]}{r}'] = f'=IF(AND({v}=1,{raw(x("别名"))}),{_alias(x("别名"))},"")'
        ws[f'{P["类型"]}{r}'] = f'=IF({v}=1,TRIM({x("类型")}&""),"")'
    # —— 规则（用户顺序）——
    G = LIST_RULE
    for r in range(2, N_RULE + 2):
        x = lambda f: _bx('规则', f)
        c = lambda f: f'{G[f]}{r}'
        kwt = 'TRIM(' + x('关键词') + '&"")'
        ws[c('词')] = '=' + esc(kwt)
        w = f'TRIM({x("查找位置")}&"")'
        ws[c('位')] = f'=IF({w}="",0,IFERROR(MATCH({w},{{"{RULE_WHERE[0]}","{RULE_WHERE[1]}"}},0),-1))'
        d = f'TRIM({x("方向")}&"")'
        ws[c('向')] = (f'=IF(OR({d}="",{d}="{RULE_DIRS[2]}"),0,IF({d}="{RULE_DIRS[0]}",1,IF({d}="{RULE_DIRS[1]}",2,-1)))')
        s = x('适用范围')
        ws[c('范围')] = (f'=IF(NOT({raw(s)}),0,IF({name2no(s, "公司")}>0,{name2no(s, "公司")},'
                         f'IF({name2no(s, "账户")}>0,100+{name2no(s, "账户")},-1)))')
        it = x('收费项')
        ws[c('项号')] = f'=IF(NOT({raw(it)}),0,IF({name2no(it, "项")}>0,{name2no(it, "项")},-1))'
        pt = x('收付人')
        ws[c('人号')] = f'=IF(NOT({raw(pt)}),0,IF({name2no(pt, "人")}>0,{name2no(pt, "人")},-1))'
        ws[c('有效')] = (f'=IF(AND({c("词")}<>"",{c("位")}>=0,{c("向")}>=0,{c("范围")}>=0,{c("项号")}>=0,{c("人号")}>=0,'
                         f'{c("项号")}+{c("人号")}>0),1,0)')
        one_i = f'IF(AND({c("有效")}=1,{c("项号")}>0),1,0)'
        one_p = f'IF(AND({c("有效")}=1,{c("人号")}>0),1,0)'
        ws[c('项计数')] = f'={one_i}' if r == 2 else f'={G["项计数"]}{r - 1}+{one_i}'
        ws[c('人计数')] = f'={one_p}' if r == 2 else f'={G["人计数"]}{r - 1}+{one_p}'
    # —— 压紧倒序：规项（第 j 行＝第 (项规则数+1−j) 条项规则）、规人 ——
    for tbl, cnt, par, n, res in ((LIST_RULE_I, '规_项计数', 'P_项规则数', N_RULE, '项号'),
                                  (LIST_RULE_P, '规_人计数', 'P_人规则数', N_RULE_P, '人号')):
        for r in range(2, n + 2):
            src = f'{tbl["源"]}{r}'
            ws[src] = f'=IF(ROW()-1>{par},0,IFERROR(MATCH({par}+2-ROW(),{cnt},0),0))'
            for f in ('词', '位', '向', '范围', res):
                ws[f'{tbl[f]}{r}'] = f'=IF({src}=0,"",INDEX(规_{f},{src}))'
    # —— 规Q：能在「收付人」里找的项规则（查找位置 收付人/空）压紧倒序（计数列对齐「规」；其余 N_RULE_Q 行；条数 P_Q规则数）——
    #    规则推断拆成「摘要那一半（扫全部项规则）＋收付人那一半（只扫这张小表）」，取用户清单里靠上的那条
    Q = LIST_RULE_Q
    for r in range(2, N_RULE + 2):
        one = f'IF(AND({G["有效"]}{r}=1,{G["项号"]}{r}>0,{G["位"]}{r}<>1),1,0)'
        ws[f'{Q["计数"]}{r}'] = f'={one}' if r == 2 else f'={Q["计数"]}{r - 1}+{one}'
    for r in range(2, N_RULE_Q + 2):
        src = f'{Q["源"]}{r}'
        ws[src] = '=IF(ROW()-1>P_Q规则数,0,IFERROR(MATCH(P_Q规则数+2-ROW(),规Q_计数,0),0))'
        for k in ('词', '向', '范围', '项号'):
            ws[f'{Q[k]}{r}'] = f'=IF({src}=0,"",INDEX(规_{k},{src}))'
    # —— 下拉压紧（第 1 行＝个数）——
    for k, pre, n in (('账户', '账户', N_ACC), ('公司', '公司', N_CO), ('收费项', '项', N_ITEM), ('收付人', '人', N_PARTY)):
        out, cnt = DV_COLS[k], DV_CNT[k]
        for r in range(2, n + 2):
            one = f'IF(INDEX({pre}_有效,ROW()-1)=1,1,0)'
            ws[f'{cnt}{r}'] = f'={one}' if r == 2 else f'={cnt}{r - 1}+{one}'
            ws[f'{out}{r}'] = f'=IF(ROW()-1>${out}$1,"",INDEX({pre}_名称,MATCH(ROW()-1,${cnt}$2:${cnt}${n + 1},0)))'
        ws[f'{out}1'] = f'={cnt}{n + 1}'
        ws[f'{cnt}1'] = f'下拉计数·{k}'
    # —— 收支分类候选键（属性 p 一列）——
    i0, g0, u0 = CAND_ROWS['项首'], CAND_ROWS['组首'], CAND_ROWS['认不出']
    for p in ATTR_NUMS:
        c = _TAB['分类候选'][f'p{p}']
        key = f'项_序键{p}'
        ws[f'{c}{CAND_ROWS["段首"]}'] = p * 10 ** 6
        for r in range(i0, i0 + N_ITEM):
            i = f'ROW()-{i0 - 1}'
            ws[f'{c}{r}'] = f'=IF(INDEX({key},{i})="","",{p * 10 ** 6}+INDEX({key},{i}))'
        for r in range(g0, g0 + N_ITEM):
            i = f'ROW()-{g0 - 1}'
            ws[f'{c}{r}'] = (f'=IF(INDEX({key},{i})="","",IF(COUNTIFS({key},">="&INDEX(项_大类号,{i})*1000,{key},"<"&INDEX({key},{i}))=0,'
                             f'{p * 10 ** 6}+INDEX(项_大类号,{i})*1000,""))')
        ws[f'{c}{u0}'] = p * 10 ** 6 + 999999
    # —— 公司→账户两级候选键 ——
    c = CAND_TWO
    for r in range(2, N_CO + 2):
        ws[f'{c}{r}'] = f'=IF(COUNTIFS(账户_公司号,ROW()-1,账户_状态号,1)>0,(ROW()-1)*100,"")'
    a0 = N_CO + 2
    for r in range(a0, a0 + N_ACC):
        a = f'(ROW()-{a0 - 1})'
        ws[f'{c}{r}'] = (f'=IF(INDEX(账户_有效,{a})<>1,"",IF(INDEX(账户_状态号,{a})>1,9900+{a},'
                         f'IF(INDEX(账户_公司号,{a})>0,INDEX(账户_公司号,{a})*100+{a},9800+{a})))')
    ws[f'{c}{a0 + N_ACC}'] = '=IF(COUNTIFS(账户_有效,1,账户_状态号,1,账户_公司号,0)>0,9800,"")'
    ws[f'{c}{a0 + N_ACC + 1}'] = '=IF(COUNTIFS(账户_有效,1,账户_状态号,">1")>0,9900,"")'
    assert a0 + N_ACC + 1 == CAND_TWO_LAST
    # —— change 10/14：月份选择格的下拉清单（名称 月份_清单＝第 2..13 行；P_管理年度＝基础资料①填的，空＝最后一笔的年份）——
    ws[f'{MONTH_LIST}1'] = '月份清单（下拉）'
    for r in range(2, 14):
        ws[f'{MONTH_LIST}{r}'] = MONTH_LIST_F


# ═════════════════════════════ 规则推断（_流 项推/人号/提本 和自检格共用一套写法） ═════════════════════════════
# LOOKUP(1,0/(条件),结果)：数组表达式里不写 IF（只用乘、加、ISNUMBER(SEARCH)）；挑写法的 IF 都在数组外面（单值）。
# R＝压紧倒序的全部项规则（v3.1 用规则区间名称 规项R_x＝INDEX(规项_x,1):INDEX(规项_x,P_项规则数)），Q＝规QR_x（能在收付人里找的）
def _tail(R, dirno, cono, accno):
    d, sc = R('向'), R('范围')
    return f'*(({d}=0)+({d}={dirno})>0)*(({sc}=0)+({sc}={cono})+({sc}=100+{accno})>0)'


def _memo_cond(R, memo, dirno, cono, accno):
    return f'ISNUMBER(SEARCH({R("词")},{memo}))*({R("位")}<>2)' + _tail(R, dirno, cono, accno)


def _party_cond(Q, party, dirno, cono, accno):
    return f'ISNUMBER(SEARCH({Q("词")},{party}))' + _tail(Q, dirno, cono, accno)


def item_lookup(R, Q, qn, user_ino, memo, party, dirno, cono, accno):
    """收费项规则推断：用户清单里最上面那条命中的项号；没有＝0。user_ino＝用户顺序的项号列（INDEX(user_ino, 第 m 条)）"""
    a = f'IF({raw(memo)},IFERROR(LOOKUP(1,0/({_memo_cond(R, memo, dirno, cono, accno)}),{R("源")}),10^6),10^6)'
    b = f'IF(AND({raw(party)},{qn}>0),IFERROR(LOOKUP(1,0/({_party_cond(Q, party, dirno, cono, accno)}),{Q("源")}),10^6),10^6)'
    return f'IFERROR(INDEX({user_ino},MIN({a},{b})),0)'


def item_multi(R, Q, qn, memo, party, dirno, cono, accno, ino):
    """「摘要能对上好几种收费项」：还有别的命中规则的项号≠ino → 那个项号（>0）；否则 0"""
    a = (f'IF({raw(memo)},IFERROR(LOOKUP(1,0/({_memo_cond(R, memo, dirno, cono, accno)}*({R("项号")}<>{ino})),'
         f'{R("项号")}),0),0)')
    b = (f'IF(AND({raw(party)},{qn}>0),IFERROR(LOOKUP(1,0/({_party_cond(Q, party, dirno, cono, accno)}*({Q("项号")}<>{ino})),'
         f'{Q("项号")}),0),0)')
    return f'MAX({a},{b})'


RI = lambda f: f'规项R_{f}'
RQ = lambda f: f'规QR_{f}'
RP = lambda f: f'规人R_{f}'


# ═════════════════════════════ _流 ═════════════════════════════
def _blk_cols(k):
    """块表某组（各 N_ACC 列）的第 a 列"""
    return lambda a: CL(CI(FLOW_BLK[k]) + a - 1)


def _above(col, r, k, inclusive):
    """本行以前（inclusive：含本行）的区域：第 1 段写往下撑的区域 col$2:col{r}（WPS 白名单）；第 2 段起定长窗口 col{r−S}:col{r}"""
    if k == 1:
        return f'{col}${R1}:{col}{r if inclusive else r - 1}'
    return win_row(col, r, inclusive=inclusive)


def _prev_late(r, k):
    """第 r 行上面所有正常行里最晚的日期＝MAX(本块「前最晚」那一格, 窗口里各行正常日)。
    窗口里上一块的行都 ≤ 本块「前最晚」那一格，多看几行结果一样（设计 §15.1.7）"""
    cell = f'${FLOW_BLK["前最晚"]}${k + 1}'
    if r == R1:
        return cell
    return f'MAX({cell},MAX({_above(FLOW["正常日"], r, k, False)}))'


def _next_day(r):
    """往下 JUMP_LOOK 个位置里第一个「有日期、跟本行日期不一样」的原始日期（空位、同一天拆的几行跳过）；没有＝0"""
    B = FLOW['日期']
    rng = f'{B}{r + 1}:{B}{r + JUMP_LOOK}'
    return f'IFERROR(INDEX({rng},MATCH(1,INDEX(({rng}>0)*({rng}<>{B}{r}),0),0)),0)'


def flow_row(r):
    """_流 第 r 行各列的公式 {字段: 公式}（AL 段号是常数）"""
    F = FLOW
    k, j, bs, _be = slot_of_row(r)
    A = f'$A{r}'
    c = lambda f: f'{F[f]}{r}'
    x = lambda f: sref(f, r)                     # INDEX(台k_字段,$A{r})：只在 $A>0 的分支里用
    tx = lambda f: f'TRIM({x(f)}&"")'
    B_, C_, D_, E_, F_, G_, H_, I_, J_, K_, L_, M_, N_, O_ = (c(f) for f in (
        '日期', '账户号', '收', '支', '有效', '对方号', '归属号', '项号', '人号', '属性号', '月码', '净', '正常日', '余额'))
    V_, Y_ = c('账日键'), c('转账键')
    AD_, AE_, AF_, AG_, AH_, AI_, AJ_, AK_, AB_ = (c(f) for f in ('日期原', '收值', '支值', '文字键', '项推', '顺序', '问本', '提本', '提醒'))
    o = {}
    # ── 取数号：本段读数区第几格（段_偏移＋j）；9 个录入格连起来 TRIM 后是空＝0（错误值算有东西，跟 段_有字 同一判法）──
    jj = f'(ROW()-{bs - 1})'
    idx = f'{seg_cell("偏移", k)}+{jj}'
    cat = '&'.join(f'INDEX({seg_name(f, k)},{idx})' for f in GATE)
    o['取数号'] = f'=IF(AND({jj}<={seg_cell("读数", k)},{idx}>=1),IF(ISERROR(1/(LEN(TRIM({cat}))=0)),{idx},0),0)'
    # ── 反查：台帐第 n 条（n＝本行行号−1）在 _流 的位置；只看段表的链 ──
    t = '(ROW()-1+INDEX(段_首,1))'
    kk = f'MATCH({t}-0.5,段_键,1)'
    o['反查'] = (f'=IFERROR(IF(OR({t}-INDEX(段_前末,{kk})>INDEX(段_读数,{kk}),INDEX(段_偏移,{kk})+{t}-INDEX(段_前末,{kk})<1),-1,'
                f'({kk}-1)*{S}+{t}-INDEX(段_前末,{kk})),-1)')
    # ── 原始列（只看本行、本段名称、_表、PAR_SAFE）──
    v = x('日期')
    tv = f'TRIM({v})'
    y8 = _ymd8(f'(--{tv})')
    o['日期原'] = (f'=IF({A}=0,"",IFERROR(IF(ISNUMBER({v}),{v},IF(AND(LEN({tv})=8,ISNUMBER(--{tv})),'
                 f'IF({y8}>0,{y8},TRIM({v}&"")),TRIM({v}&""))),"#错"))')
    d = AD_
    s_ = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(LEFT({d},10),".","-"),"/","-"),"年","-"),"月","-"),"日","")'
    y4 = f'IFERROR(AND(--LEFT({d},4)>=2000,--LEFT({d},4)<=2099),FALSE)'
    o['日期'] = (f'=IF({A}=0,0,IF(ISNUMBER({d}),IF(AND({d}>=36526,{d}<73051),INT({d}),{_ymd8(d)}),'
               f'IF({y4},IFERROR(IF(AND(--({s_})>=36526,--({s_})<73051),INT(--({s_})),0),0),0)))')

    def numv2(xx):
        tt = f'TRIM({xx}&"")'
        ss = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({tt},",",""),"，",""),"¥",""),"￥","")'
        return (f'IF(ISNUMBER({xx}),ROUND({xx},2),IF(OR({tt}="",{tt}="-",{tt}="--",{tt}="—",{tt}="－"),0,'
                f'IFERROR(ROUND(--{ss},2),IF(ISTEXT({xx}),"坏",0))))')
    o['收值'] = f'=IF({A}=0,0,IFERROR({numv2(x("收入"))},"坏"))'
    o['支值'] = f'=IF({A}=0,0,IFERROR({numv2(x("支出"))},"坏"))'
    # 截到 200 字：重复粘贴用它当 COUNTIFS 的条件，Excel 的条件超过 255 字就 #VALUE!（v2 是三格各比各的）
    o['文字键'] = f'=IF({A}=0,"",LEFT({tx("摘要")}&"|"&{tx("备注")}&"|"&{tx("原表位置")},200))'
    # 账户格是错误值（#N/A 这种）＝0（不往块表传错：块表一错，后面整本的余额、正常日都错）；问本第一条报「有错误值」
    o['账户号'] = f'=IF({A}=0,0,IFERROR({name2no(x("账户"), "账户")},0))'
    o['收'] = f'=ROUND(MAX(N({AE_}),0)+MAX(-N({AF_}),0),2)'
    o['支'] = f'=ROUND(MAX(N({AF_}),0)+MAX(-N({AE_}),0),2)'
    o['有效'] = f'=IF(AND({B_}>=P_开账日,{B_}<=P_日期上限,{C_}>0,OR({D_}<>0,{E_}<>0)),1,0)'
    o['对方号'] = f'=IF({A}=0,0,IFERROR({name2no(x("对方账户"), "账户")},0))'
    xc = x('归属公司')
    o['归属号'] = (f'=IF({C_}=0,0,IF(TRIM({xc}&"")="",INDEX(账户_公司号,{C_}),'
                 f'IFERROR(MATCH({esc(norm(xc))},公司_名称规,0),INDEX(账户_公司号,{C_}))))')
    memo, party = x('摘要'), x('收付人')
    dirno, cono = f'(2-({D_}>0))', f'INDEX(账户_公司号,{C_})'
    internal = f'AND({G_}>0,{G_}<>{C_},P_内部号>0)'
    # 收费项填了（导入的数据全是）就不碰摘要、收付人（IF 挡住，省得每行都算一遍 raw）
    o['项推'] = (f'=IF({A}=0,0,IF(OR({raw(x("收费项"))},{internal}),0,IF(AND({C_}>0,P_项规则数>0,OR({raw(memo)},{raw(party)})),'
               f'{item_lookup(RI, RQ, "P_Q规则数", "规_项号", memo, party, dirno, cono, C_)},0)))')
    o['项号'] = (f'=IF({A}=0,0,IF({raw(x("收费项"))},{name2no(x("收费项"), "项")},'
               f'IF({internal},P_内部号,{AH_})))')
    cond_p = f'ISNUMBER(SEARCH({RP("词")},{memo}))*({RP("位")}<>2)' + _tail(RP, dirno, cono, C_)
    o['人号'] = (f'=IF({A}=0,0,IF({raw(party)},{name2no(party, "人")},'
               f'IF(AND({C_}>0,P_人规则数>0,{raw(memo)}),IFERROR(LOOKUP(1,0/({cond_p}),{RP("人号")}),0),0)))')
    unk = f'IF({D_}>0,1,2)' if ATTR_UNKNOWN == 0 else str(ATTR_UNKNOWN)
    o['属性号'] = (f'=IF({I_}>0,IF(INDEX(项_属性号,{I_})>0,INDEX(项_属性号,{I_}),IF({D_}>0,1,2)),'
                f'IF(AND({G_}>0,{G_}<>{C_}),{INT9},{unk}))')
    o['月码'] = f'=IF({F_}=1,YEAR({B_})*100+MONTH({B_}),0)'
    o['净'] = f'=IF({F_}=1,ROUND({D_}-{E_},2),0)'
    # ── 挂在块表累计链上的：正常日、余额、顺序（直接写本块那一行的块表格；第 2 段起定长窗口）──
    prev = _prev_late(r, k)
    nxt = _next_day(r)
    # 跳期：比 prev 晚 JUMP_DAYS 天以上，而且下一笔早回去 JUMP_DAYS 天以上；下一笔＝0（后面没有）时也算（△）
    o['正常日'] = (f'=IF({F_}<>1,0,IF({B_}>{prev}+{JUMP_DAYS},IF(AND({prev}>=P_开账日,{B_}>{nxt}+{JUMP_DAYS}),0,{B_}),{B_}))')
    if k == 1:
        ssum = f'SUMIFS({_above(F["净"], r, 1, True)},{_above(F["账户号"], r, 1, True)},{C_})'
    else:
        ssum = (f'SUMIFS({win_row(F["净"], r)},{win_row(F["账户号"], r)},{C_},'
                f'{win_row(F["段号"], r)},{F["段号"]}{r})')
    o['余额'] = f'=IF({F_}<>1,"",ROUND(INDEX(账户_期初,{C_})+INDEX({blk_row("前", k)},{C_})+{ssum},2))'
    late = f'INDEX({blk_row("账前最晚", k)},{C_})'
    if r == R1:
        order = f'{B_}<{late}'
    else:
        kb = _above(F['账日键'], r, k, False)
        order = f'IF({B_}<{late},1,COUNTIFS({kb},">"&{V_},{kb},"<"&({C_}+1)*{K_DAY}))>0'
    jump_x = (f'"※ 比前面最晚的一笔（"&MONTH({prev})&"/"&DAY({prev})&"）晚 "&({B_}-{prev})&" 天，下一笔又早回去：'
              f'年份、月份写错了？"')
    jump_w = f'"△ 比前面最晚的一笔晚 "&({B_}-{prev})&" 天，日期没写错吧？"'
    o['顺序'] = (f'=IF({F_}<>1,"",IF({N_}=0,IF({nxt}>0,{jump_x},{jump_w}),'
               f'IF({order},"{T_ORDER}",IF({O_}<0,"{T_NEG}",""))))')
    # ── 问本：※ 本行链（不含跳期）──
    xin_bad = f'OR(ISTEXT({AE_}),ISTEXT({AF_}))'
    allf = '&'.join(x(f) for f in LEDGER_COLS)
    pairs = [
        (f'ISERROR({allf})', '"※ 这一行有错误值（#N/A、#VALUE! 这种）：改成正常的日期、文字、数字"'),
        (f'{d}=""', '"※ 没填日期"'),
        (f'{B_}=0', '"※ 日期看不懂：要写 2026/9/1 这样的日期"'),
        (f'{B_}<P_开账日', '"※ 早于开账日：上一年的账记在上一本"'),
        (f'{B_}>P_日期上限', '"※ 日期晚于 "&YEAR(P_日期上限)&" 年底：年份写错了？"'),
        (f'NOT({raw(x("账户"))})', '"※ 没选账户"'),
        (f'{C_}=0', f'"※ 账户「"&{tx("账户")}&"」不在基础资料③：改了名就把旧名加进别名"'),
        (xin_bad, '"※ 金额看不懂：10000+200 要写成 =10000+200"'),
        (f'AND({D_}=0,{E_}=0)', '"※ 没填金额"'),
        (f'AND(N({AE_})<>0,N({AF_})<>0)', '"※ 收入、支出都填了：分成两行记"'),
        (f'IF({raw(x("收费项"))},{I_}=0,FALSE)', f'"※ 收费项「"&{tx("收费项")}&"」不在④：从下拉选，或去④加（改了名就加别名）"'),
        (f'AND({K_}={INT9},NOT({raw(x("对方账户"))}))', '"※ 内部转账要选转账的另一户；对方不在本册就改成往来款、借款等"'),
        (f'IF({raw(x("对方账户"))},{G_}=0,FALSE)', f'"※ 转账的另一户「"&{tx("对方账户")}&"」不在③：外单位、个人请填在收付人"'),
        (f'AND({G_}>0,{G_}={C_})', '"※ 转账的另一户就是本账户"'),
        (f'AND({G_}>0,{K_}<>{INT9})', '"※ 填了转账的另一户，收费项要选内部转账（原性质写摘要）"'),
        (f'IF({raw(xc)},{name2no(xc, "公司")}=0,FALSE)', f'"※ 归属公司「"&{tx("归属公司")}&"」不在②"'),
    ]
    o['问本'] = f'=IF({A}=0,"",{chain(pairs)})'
    # ── 提本：△ 本行链（只看本行和前后 DUP_WIN 个位置，不碰块表）──
    lo, hi = max(R1, r - DUP_WIN), r + DUP_WIN
    W = lambda f: f'{F[f]}{lo}:{F[f]}{hi}'
    num = f'{W("账日键")},{V_},{W("收")},{D_},{W("支")},{E_}'
    dup = f'IF(COUNTIFS({num})>1,IFERROR(COUNTIFS({num},{W("文字键")},"="&{esc(AG_)}),0),0)>1'
    guessed = f'AND({AH_}>0,{G_}=0)'
    multi = item_multi(RI, RQ, 'P_Q规则数', memo, party, dirno, cono, C_, I_)
    iname = f'INDEX(项_名称,{I_})'
    textc = f'OR(ISTEXT({x("日期")}),AND(ISTEXT({x("收入")}),{raw(x("收入"))}),AND(ISTEXT({x("支出")}),{raw(x("支出"))}))'
    me = f'{x("摘要")}&{x("备注")}'
    refund = f'OR(ISNUMBER(SEARCH("退",{me})),ISNUMBER(SEARCH("冲",{me})),ISNUMBER(SEARCH("返还",{me})))'
    paused = f'AND(INDEX(账户_状态号,{C_})>1,OR(INDEX(账户_停用日,{C_})=0,{B_}>INDEX(账户_停用日,{C_})))'
    t2 = [
        (dup, f'"{T_DUP}"'),
        (f'AND({I_}>0,{I_}=P_待定号)', '"△ 收费项是「待定」：定了性质请改过来"'),
        (guessed, f'IF({multi}>0,"△ 摘要能对上好几种收费项，按「"&{iname}&"」算了：请在收费项里选",'
                  f'"△ 收费项按摘要认成「"&{iname}&"」：对就在收费项选上")'),
        (f'AND(NOT({raw(x("收费项"))}),{I_}=0)', '"△ 收费项没填，摘要也认不出：请选收费项"'),
        (textc, '"△ 日期或金额是文字，已照认：最好改成真日期、数字"'),
        (f'OR(N({AE_})<0,N({AF_})<0)', '"△ 金额是负数，已当另一栏算"'),
        (f'IF(OR(AND({K_}=1,{E_}>0),AND({K_}=2,{D_}>0)),NOT({refund}),FALSE)', f'IF({K_}=1,"{T_REFUND_IN}","{T_REFUND_OUT}")'),
        (f'IF({raw(party)},{J_}=0,FALSE)', f'"△ 收付人「"&{tx("收付人")}&"」不在⑤：报表算没填"'),
        (f'AND({J_}=0,OR({K_}=1,{K_}=2))', f'"{T_NOPARTY}"'),
        (paused, '"△ 账户已暂停/销户，还在记账"'),
    ]
    o['提本'] = f'=IF({F_}<>1,"",{chain(t2)})'
    # ── 短合并列（只引用本行）──
    o['提醒'] = f'=IF({F_}<>1,"",IF(LEFT({AI_},1)="△",{AI_},{AK_}))'
    o['问题'] = f'=IF({A}=0,"",IF({AJ_}<>"",{AJ_},IF(LEFT({AI_},1)="※",{AI_},{AB_})))'
    # ── 键 ──
    yr, mo = f'YEAR({B_})', f'MONTH({B_})'
    o['账年属月键'] = f'=IF({F_}=1,{k_aypm(C_, yr, K_, mo)},0)'
    o['归年属月键'] = f'=IF({F_}=1,{k_aypm(H_, yr, K_, mo)},0)'
    o['项年属月键'] = f'=IF({F_}=1,{k_aypm(I_, yr, K_, mo)},0)'
    o['人年属月键'] = f'=IF({F_}=1,{k_aypm(J_, yr, K_, mo)},0)'
    o['项归年属月键'] = f'=IF({F_}=1,{k_igypm(I_, H_, yr, K_, mo)},0)'
    o['账日键'] = f'=IF({F_}=1,{C_}*{K_DAY}+{B_},0)'
    o['账年键'] = f'=IF({F_}=1,{C_}*10000+{yr},0)'
    o['账月键'] = f'=IF({F_}=1,{C_}*{K_YM}+{L_},0)'
    o['转账键'] = (f'=IF(AND({F_}=1,{K_}={INT9},{G_}>0,{G_}<>{C_}),'
                f'(MIN({C_},{G_})*{K_PAIR}+MAX({C_},{G_}))*{K_DAY}+{B_},0)')
    # ── 录入模式停算（FLOW_GATED）──
    lo_ = win_first(ym_of(f'({B_}-{PAIR_DAYS})'))
    hi_ = win_last(ym_of(f'({B_}+{PAIR_DAYS})'), lo_)
    wy, wm = win('流_转账键', lo_, hi_), win('流_净', lo_, hi_)
    o['配对'] = pause(f'IF({Y_}=0,"",IF(ROUND(SUMIF({wy},{Y_},{wm}),2)=0,0,'
                    f'IF(COUNTIFS({wy},">="&({Y_}-{PAIR_DAYS}),{wy},"<="&({Y_}+{PAIR_DAYS}),{wm},-{M_})'
                    f'>=COUNTIFS({wy},">="&({Y_}-{PAIR_DAYS}),{wy},"<="&({Y_}+{PAIR_DAYS}),{wm},{M_}),0,1)))')
    o['期人键'] = pause0(f'IF(AND({F_}=1,{B_}>=P_期起,{B_}<=P_期止),{J_}*10+{K_},0)')
    return o


def build_seg(ws):
    """段表（GY:HM，第 k 段在第 k+1 行；名称 段_xx；公式原文见设计 §15.1.3）"""
    G = FLOW_SEG
    for key, col in G.items():
        ws[f'{col}1'] = f'段·{key}'
    for k in range(1, N_SEG + 1):
        r = k + 1
        c = lambda key: f'{G[key]}{r}'
        up = lambda key: f'{G[key]}{r - 1}'
        U, W = seg_area('锚', k), read_area('锚', k)
        ws[c('首')] = f'=ROW(INDEX({U},1))'
        ws[c('末')] = f'=ROW(INDEX({U},ROWS({U})))'
        ws[c('区首')] = f'=IFERROR(ROW(INDEX({W},1)),10^6)'
        ws[c('区末')] = f'=IFERROR(ROW(INDEX({W},ROWS({W}))),0)'
        ws[c('尾')] = f'=IFERROR({c("末")},{c("首")})' if k == 1 else f'=MAX({up("尾")},IFERROR({c("末")},{up("尾")}))'
        ws[c('前末')] = f'={c("首")}' if k == 1 else f'={up("读到")}'
        ws[c('读到')] = f'=MAX({c("前末")},MIN({c("前末")}+{S},{c("尾")},{c("区末")}))'
        ws[c('读数')] = f'={c("读到")}-{c("前末")}'
        ws[c('偏移')] = f'={c("前末")}-{c("区首")}+1'
        ws[c('键')] = f'={c("前末")}+{k}/1000'
        ws[c('漏行')] = f'=MAX(0,MIN({c("读到")},{c("区首")}-1)-{c("前末")})'
        prevtail = c('首') if k == 1 else up('尾')
        ws[c('夹缝')] = f'=MAX(0,MIN({c("尾")},{c("区首")}-1)-{prevtail})'
        if k < N_SEG:
            ws[c('状态')] = f'=IF({c("漏行")}>0,2,IF(ISERROR({c("末")}),3,IF({c("尾")}>{c("读到")},1,0)))'
        else:
            ws[c('状态')] = f'=IF({c("漏行")}>0,2,IF(ISERROR({c("末")}),3,0))'
        # 叶子：错位＝录入列名称跟锚列读数区（首行、行数）不一样的个数；锚列读数区出错时：数据名称也出错＝0，否则 1
        terms = '+'.join(f'IFERROR((ROW(INDEX({seg_name(f, k)},1))<>{c("区首")})+(ROWS({seg_name(f, k)})<>{c("区末")}-{c("区首")}+1),1)'
                         for f in LEDGER_COLS)
        ws[c('错位')] = (f'=IF(ISERROR(ROW(INDEX({W},1))),IF(ISERROR(ROW(INDEX({seg_name("日期", k)},1))),0,1),{terms})')
        # 叶子：有字＝本段归属又在读数区里的行（最后一段多看到 读到+S）9 个录入格有东西的行数；错误值逐列另数（LO 数组 & 吞错误值）
        s_ = f'MAX(1,{prevtail}-{c("区首")}+2)'
        e_ = (f'{c("尾")}-{c("区首")}+1' if k < N_SEG else f'MIN({c("尾")},{c("读到")}+{S})-{c("区首")}+1')
        rng = lambda f: f'INDEX({seg_name(f, k)},{s_}):INDEX({seg_name(f, k)},{e_})'
        cat = '&'.join(rng(f) for f in GATE)
        errs = '+'.join(f'ISERROR({rng(f)})' for f in GATE)
        ws[c('有字')] = f'=IFERROR(IF({e_}<{s_},0,SUMPRODUCT(--(ISERROR(1/(LEN(TRIM({cat}))=0))+{errs}>0))),0)'


def build_flow(ws, ctx):
    for c, f, _d in FLOW_COLS:
        ws[f'{c}1'] = f
    seg_col = FLOW['段号']
    for r in range(R1, RN + 1):
        o = flow_row(r)
        for f, v in o.items():
            ws[f'{FLOW[f]}{r}'] = v
        ws[f'{seg_col}{r}'] = slot_of_row(r)[0]            # 段号：常数（FLOW_CONST）
    assert set(FLOW_CONST) == {'段号'}
    # ── 块表（一块＝一段）──
    net, pre, late, plate = (_blk_cols(k) for k in ('净', '前', '账最晚', '账前最晚'))
    for k, t in (('净', '块内净额'), ('前', '块前累计'), ('账最晚', '块内最晚日'), ('账前最晚', '块前最晚日')):
        for a in range(1, N_ACC + 1):
            ws[f'{_blk_cols(k)(a)}1'] = f'{t}·{a}'
    for k in ('最晚', '前最晚', '最晚月', '至最晚月', '最早月', '后最早月'):
        ws[f'{FLOW_BLK[k]}1'] = f'块·{k}'
    Cc, Mc, Nc, Lc, Vc = (FLOW[k] for k in ('账户号', '净', '正常日', '月码', '账日键'))
    for b in range(1, N_BLK + 1):
        r = b + 1
        bs, be = (b - 1) * BLK + 2, b * BLK + 1
        rg = lambda col: f'${col}${bs}:${col}${be}'
        for a in range(1, N_ACC + 1):
            ws[f'{net(a)}{r}'] = f'=SUMIFS({rg(Mc)},{rg(Cc)},{a})'
            ws[f'{pre(a)}{r}'] = '=0' if b == 1 else f'={pre(a)}{r - 1}+{net(a)}{r - 1}'
            ws[f'{late(a)}{r}'] = (f'=MAX(0,IFERROR(LARGE({rg(Vc)},COUNTIF({rg(Vc)},">="&{(a + 1) * K_DAY})+1)-{a * K_DAY},0))')
            ws[f'{plate(a)}{r}'] = '=0' if b == 1 else f'=MAX({plate(a)}{r - 1},{late(a)}{r - 1})'
        cl = lambda k: FLOW_BLK[k]
        ws[f'{cl("最晚")}{r}'] = f'=MAX({rg(Nc)})'
        ws[f'{cl("前最晚")}{r}'] = '=0' if b == 1 else f'=MAX({cl("前最晚")}{r - 1},{cl("最晚")}{r - 1})'
        ws[f'{cl("最晚月")}{r}'] = f'=MAX({rg(Lc)})'
        ws[f'{cl("至最晚月")}{r}'] = (f'={cl("最晚月")}{r}' if b == 1 else f'=MAX({cl("至最晚月")}{r - 1},{cl("最晚月")}{r})')
        ws[f'{cl("最早月")}{r}'] = f'=IFERROR(SMALL({rg(Lc)},COUNTIF({rg(Lc)},0)+1),999999)'
        ws[f'{cl("后最早月")}{r}'] = (f'={cl("最早月")}{r}' if b == N_BLK else f'=MIN({cl("最早月")}{r},{cl("后最早月")}{r + 1})')
    build_seg(ws)
    _probe(ws)


# —— 规则推断自检（设计 §1.5.4 perf#8）：常数规则 6 条（倒序存，跟压紧表一样）＋3 个同样写法的 LOOKUP 格 ——
PROBE_RULES = [  # 用户顺序：词（esc 前）, 位, 向, 范围, 项号
    ('手续费', 1, 2, 0, 11), ('A*B', 0, 0, 0, 12), ('租金', 0, 1, 3, 13), ('顺丰', 2, 0, 0, 14), ('费', 1, 0, 105, 15), ('租', 0, 0, 0, 16)]
PROBE_CASES = [  # 摘要, 收付人, 方向号, 公司号, 账户号, 期望（用户清单里最上面那条命中的项号）
    ('A105 3月租金', '', 1, 3, 7, 13), ('AxB 租金手续费', '顺丰速运', 1, 2, 5, 14), ('退场地费', '', 2, 1, 5, 15)]
PROBE_HDR = ['词', '位', '向', '范围', '项号', '源', '用户项号', 'Q源', 'Q词', 'Q向', 'Q范围', 'Q项号',
             '摘要', '收付人', '方向号', '公司号', '账户号', '期望', '结果', '不一致']


def probe_col(k):
    return CL(CI(FLOW_PROBE['首列']) + PROBE_HDR.index(k))


def _probe(ws):
    """跟 _流 项推一模一样的写法（item_lookup：摘要一半＋收付人一半取靠上的）算 3 个例子，期望值生成时写死"""
    n = len(PROBE_RULES)
    assert n == FLOW_PROBE['行数']
    for k in PROBE_HDR:
        ws[f'{probe_col(k)}1'] = '自检·' + k
    qrules = [k for k in range(n, 0, -1) if PROBE_RULES[k - 1][1] != 1]      # 能在收付人里找的（倒序）
    for j in range(1, n + 1):                   # 第 j 行＝第 (n+1−j) 条（倒序）
        src = n + 1 - j
        w, pos, d, sc, it = PROBE_RULES[src - 1]
        r = j + 1
        ws[f'{probe_col("词")}{r}'] = '=' + esc(f'"{w}"')
        ws[f'{probe_col("位")}{r}'], ws[f'{probe_col("向")}{r}'] = pos, d
        ws[f'{probe_col("范围")}{r}'], ws[f'{probe_col("项号")}{r}'] = sc, it
        ws[f'{probe_col("源")}{r}'] = src
        ws[f'{probe_col("用户项号")}{r}'] = PROBE_RULES[j - 1][4]
    for j, src in enumerate(qrules, 1):
        w, pos, d, sc, it = PROBE_RULES[src - 1]
        r = j + 1
        ws[f'{probe_col("Q源")}{r}'] = src
        ws[f'{probe_col("Q词")}{r}'] = '=' + esc(f'"{w}"')
        ws[f'{probe_col("Q向")}{r}'], ws[f'{probe_col("Q范围")}{r}'], ws[f'{probe_col("Q项号")}{r}'] = d, sc, it
    rg = lambda k: f'${probe_col(k)}$2:${probe_col(k)}${n + 1}'
    R = lambda f: f'INDEX({rg(f)},1):INDEX({rg(f)},{n})'
    Q = lambda f: f'INDEX({rg("Q" + f)},1):INDEX({rg("Q" + f)},{len(qrules)})'
    for i, (memo, party, d, co, acc, exp) in enumerate(PROBE_CASES):
        r = i + 2
        cc = lambda k: f'{probe_col(k)}{r}'
        ws[cc('摘要')], ws[cc('收付人')] = memo, party or None
        ws[cc('方向号')], ws[cc('公司号')], ws[cc('账户号')], ws[cc('期望')] = d, co, acc, exp
        ws[cc('结果')] = '=' + item_lookup(R, Q, str(len(qrules)), rg('用户项号'), cc('摘要'), cc('收付人'),
                                         cc('方向号'), cc('公司号'), cc('账户号'))
        ws[cc('不一致')] = f'=IF({cc("结果")}={cc("期望")},0,1)'


# ═════════════════════════════ _参 ═════════════════════════════
def _far_below():
    """容量外、离得远的行：最后一段的 段_有字 只看到 读到＋S 行；再往下还有日期或账户的行数（整列名称，表头行起算）"""
    i0 = f'P_台帐末-INDEX(段_首,1)+{S + 2}'
    cnt = lambda f: cnta(f'INDEX({tref(f)},{i0}):INDEX({tref(f)},ROWS({tref(f)}))')     # a74split：不数空串（查询本取数格）
    return f'IFERROR(IF(ISNUMBER({i0}),MAX({cnt("日期")},{cnt("账户")}),0),0)'


def build_par(ws, ctx):
    bp = lambda k: f'{q(SH_BASE)}!${BASE_PAR_VAL}${PAR_ROW[k]}'
    open_d = ctx['params'].get('开账日')
    dflt = f'DATE({open_d.year},{open_d.month},{open_d.day})' if open_d else 'DATE(2026,1,1)'
    mg = bp('管理年度')
    pst = f'TRIM({bp("录入模式")}&"")'
    pc = probe_col('不一致')
    home_sel = f'{q(SH_HOME)}!${HOME_SEL["截止"][0]}${HOME_SEL["截止"][1:]}'
    span = lambda k: f'{q(SH_SPAN)}!${SPAN_SEL[k][0]}${SPAN_SEL[k][1:]}'
    rows = {
        '集团名称': (f'=TRIM({bp("集团名称")}&"")', '基础资料①（空着：标题自己写「集团」）'),
        '开账日': ('=' + dsel(bp('开账日'), dflt), '基础资料①（dsel）'),
        '开账年': ('=YEAR(P_开账日)', ''),
        '日期上限': (f'=DATE(P_开账年+{N_YEARS - 1},12,31)', '更晚的 ※ 不算'),
        '最后日期': ('=MAX(P_开账日,MAX(块_最晚))', '有效、没跳期的流水里最晚的一天'),
        '管理年度': (f'=IF(ISNUMBER({mg}),IF(AND({mg}>=2000,{mg}<=2100),INT({mg}),YEAR(P_最后日期)),YEAR(P_最后日期))',
                  '基础资料①填了 2000～2100 用它；否则最后日期的年份'),
        '截止': ('=' + dsel(home_sel, 'IF(P_管理年度<YEAR(P_最后日期),DATE(P_管理年度,12,31),P_最后日期)'), '首页 C3'),
        '截止年': ('=YEAR(P_截止)', ''),
        '截止月码': ('=YEAR(P_截止)*100+MONTH(P_截止)', ''),
        '期起': ('=' + dsel(span('开始'), 'DATE(YEAR(P_最后日期),1,1)'), '任意时间段 B3'),
        '期止': ('=' + dsel(span('结束'), 'P_最后日期'), '任意时间段 B4'),
        '已用行数': (f'=MAX({cnta(lrange("日期"))},{cnta(lrange("账户"))})', '资金台帐 日期或账户有内容的行数（不数空串：查询本取数格空着是 ""）'),
        '期初合计': ('=SUM(账户_期初)', ''),
        '待定号': (f'=IFERROR(MATCH("{ITEM_PENDING}",项_名称规,0),0)', ''),
        '内部号': (f'=IFERROR(MATCH({INT9},项_属性号,0),0)', '第一个属性「内部转账」的收费项'),
        '账户数': ('=COUNTIF(账户_有效,1)', ''),
        '公司数': ('=COUNTIF(公司_有效,1)', ''),
        '收费项数': ('=COUNTIF(项_有效,1)', ''),
        '收付人数': ('=COUNTIF(人_有效,1)', ''),
        '规则数': ('=COUNTIF(规_有效,1)', '有效规则条数'),
        '项规则数': ('=MAX(规_项计数)', ''),
        '人规则数': (f'=MIN({N_RULE_P},MAX(规_人计数))', f'超过 {N_RULE_P} 条核对表 ※'),
        '录入模式': (f'=IF(OR({pst}="是",{pst}="1",{pst}="录入",{pst}="暂停"),1,0)', '基础资料①「录入模式」（v2 的「暂停」照认）'),
        '自检': (f'=SUM({H_FLOW}!${pc}$2:${pc}${len(PROBE_CASES) + 1})', '规则推断自检不一致的个数（0＝正常）'),
        # —— v3.1 分段（设计 §15.1.3、§15.1.5；录入模式照算）——
        '台帐末': (f'=INDEX(段_读到,{N_SEG})', '读得到的最后一行（再往下有字＝容量外）'),
        '台帐容量': ('=IFERROR(P_台帐末-INDEX(段_首,1),0)', '台帐能录的行数（表头以下到 台帐末）'),
        '漏行': ('=SUM(段_漏行)', '读不到的行数（含空行，只给核对说明定位）'),
        'Q规则数': ('=MAX(规Q_计数)', f'能在收付人里找的项规则条数（超过 {N_RULE_Q} 核对 ※）'),
        '错位': ('=IFERROR(SUM(段_错位),1)', '>0：台帐被插/删了单元格而不是整行、只剪了 A:P、删了列'),
        '笔数差': ('=SUM(段_有字)-COUNTIF(流_取数号,">0")', '有字没读到的笔数 − 读了两遍的笔数（≠0 就 ※）'),
        '没算': ('=MAX(0,P_笔数差)+IFERROR(SUM(段_夹缝),0)+' + _far_below(), '有几行没算进来（含容量外远处的）'),
        '锚坏': ('=IFERROR(IF(OR(ISERROR(INDEX(段_首,1)),COUNT(段_末)=0,NOT(ISTEXT(INDEX(台帐_日期,1)))),1,0),1)',
               '1＝台帐的锚列 Q 或表头行被删'),
        '可录': ('=MAX(0,P_台帐容量-P_已用行数)', '还能录约几行'),
    }
    assert set(rows) == set(PAR_KEYS), set(PAR_KEYS) ^ set(rows)
    ws['A1'], ws['B1'], ws['C1'] = '参数', '值', '说明'
    for k, cell in PAR.items():
        r = int(cell[1:])
        f, note = rows[k]
        ws[f'A{r}'] = k
        ws[cell] = f
        if note:
            ws[f'C{r}'] = note
    for k in ('开账日', '日期上限', '最后日期', '截止', '期起', '期止'):
        ws[PAR[k]].number_format = 'yyyy/mm/dd'
    ws.column_dimensions['A'].width = 12
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 40


def build(wb, ctx):
    build_tab(wb[H_TAB], ctx)
    build_flow(wb[H_FLOW], ctx)
    build_par(wb[H_PAR], ctx)
