# -*- coding: utf-8 -*-
"""A074 集团多账户资金台帐 · 【首页】（资金看板，老板看）和【年末余额】（设计.md §5.1、§5.9；v3.1 改动见 §15.2.2、§15.1.5）。

口径（设计 §1.3、§1.4）：一律按「日期」算，截到首页 C3 的截止日（P_截止；空＝最后一笔那天，管理年度填了上一年＝那年 12/31）。
  资金（余额）＝期初＋日期 ≤ 截止的收－支，按「账户所属公司」汇总；暂停/销户户单列；
  经营收入＝属性「收入」行 Σ净，经营支出＝属性「支出」行 −Σ净，投资建设＝−Σ净，往来流入/流出＝往来行收入栏/支出栏，
  待分类＝Σ净，内部转入/转出＝内部转账行收入栏/支出栏；公司的「本年经营净额」按「归属公司」（流_归年属月键）。
  当天＝截止那一天；本月＝截止那个月 1 号～截止；本年＝截止那年 1/1～截止（不含截止以后的流水）。
  账户行的本月收入/支出是账户口径（收入栏/支出栏，不含内部转账）。

v3.1（设计 §15）：
- 录入模式（P_录入模式＝1；首页、年末余额都在 PAUSE_SHEETS）：第 2 行起每个公式最外层 pause()/pause0()（make 查）。
  首页只留（登记 ws._pause_exempt）：第 3 行抬头（按…算、数据到、台帐已用/能录）、提醒条 ※/△ 两格、资金总额 C6（改成
  P_期初合计＋SUM(块_净)，A5 标签写「录入模式：全部流水，不按截止日」）、四个段标题（换成「录入模式中先不算」）。
  第 1 行右边 J1:R1 是横幅（第 1 行 make 不查；录入模式也算）：按设计 §15.1.5 先后只显示最要紧的一条
  锚坏 ＞ 错位/读了两遍 ＞ 没算进来 ＞ 满了/快满 ＞ 录入模式（跟台帐 F1 同一套 P_xx、同样的话）。
  年末余额全停：A1 换成 layout.PAUSE_TITLE（红字），M1（年末_年份）照算。
- 取数一律走名称、同表格子：不直接写「别的表!」（粘贴区原文用 bref('账户', 字段)），不写往下撑的区域（第几个用 ROW()−常数；首页、年末余额有保护，不插行）。
- 复核 F11：提醒条「待分类」写明是净额，点过去是台帐 M 列（属性）表头，提示「点筛选只勾待分类」（超链接的屏幕提示＋第 2 行说明）。
- 复核 F4：导航里备用子表那句改成「表名别改」（首页、核对表的链接按表名跳）。备用子表 2 张（layout.SPARE_NAMES）。
- 「台帐已用」改成「已用 x 行 / 能录 y 行（还能录约 z 行）」（P_台帐容量、P_可录）；提醒条「新本」的快满按 P_可录（同横幅）。

版面（首页，A..R；行号见下面常量，model/chk 也用）：
  1 标题 A1:I1（集团名称、本本年份、开账日）＋横幅 J1:R1｜2 说明｜3 截止日期（黄 C3）、按…算、数据到、台帐已用/能录｜4 提醒条（9 格，带跳转）
  5～8 资金总额卡片（C6＝首页_资金总额、C7 在用、C8 暂停/销户）＋当天/本月/本年 12 项（E..R，J 是隔条）
  10 段标题｜11 表头｜12～41 各公司（30 行位，A..I）＋ 右边 12～23 行 1～12 月、24 行合计（K..R）｜42 公司合计
  44 段标题｜45 表头｜46～117 公司→账户两级（72 行位，A..I）｜118 合计；右边 45 行起走势图（柱：经营收入/支出；线：月末资金，右轴）、
  59 行起导航（报表、录入、核对、说明；26 张子表按公司排、旁边小字「原：xx」；备用子表一行）。
  生成时按 ctx 的公司数、两级行数把多余的行位隐藏（公司至少留 13 行给右边的 12 个月表）；段标题会提示「下面有藏起来的行」。
私有隐藏列（F_HELP 小灰字，不进 layout.NAMES）：
  Y/Z 第 1～10 行标量（月初、月窗、年窗、余额窗末、开账月码、最后月码、跳期笔数、没记账户数）；
  AB..AO 第 2～41 行每个账户（号、两级组号、公司号、类型号、状态号、截止余额、本月收入、本月支出、最后键、最后一笔、几天没记、原名、备注、
    最晚＝这个账户所有有效流水里最晚的日期＝MAX(块_账最晚 那一列)）。v3.1 提速：最晚 ≤ 截止（这个账户没有截止以后的流水，平常都是）时，
    余额＝期初＋SUM(块_净 那一列)、最后一笔＝最晚，不再扫 _流 窗口（LARGE/SUMIFS 只在有截止以后的流水时才算；结果一样）；
  AP/AQ 第 2～27 行 子表名规 → 原表名（生成时写死，给「原：xx」用）；
  AS..AX 第 12～23 行每月（月码、窗首、窗末、显示、线窗末、图上的月末资金＝截到 MIN(月末, 截止)，未来月份是截止日余额，线不会掉到 0；
    v3.1 提速：2～12 月＝上个月＋本月（只扫本月窗口），只有 1 月扫 [1, 线窗末]）；
  BA..BC 第 46～117 行两级清单（键、账户号、组号）。
年末余额（layout YE_*）：B3 年份（黄，空＝管理年度），M1＝生效年份；6～45 行每个账户（第 k 行位＝基础资料③第 k 行），46 合计；
  49 粘贴区标题、50 开新本五步（红字）、51 表头＝基础资料③ B..M 的表头、52～91 行跟③ 一列对一列（同样是 B..M 列；期初余额＝年末余额）。
  私有：O1 年初窗末、O2 年末窗末，O6..O45 每行账户号，P6..P45 每行账户的最晚日期（MAX(块_账最晚 那一列)；
  v3.1 提速：年份 ≤ 开账年 → 年初＝期初；最晚 ≤ 那年 12/31 → 年末＝期初＋SUM(块_净 那一列)；否则照旧扫窗口）。
给后面模块用：首页 C3（_参 读 P_截止）、C6（首页_资金总额；录入模式＝期初合计＋全部流水净额）、年末余额 M1（年末_年份）、I46（年末_合计）。"""
import re
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
import data
from layout import (
    SH_HOME, SH_LEDGER, SH_BASE, SH_DAY, SH_MON, SH_CO, SH_CLS, SH_ITEMCO, SH_IE, SH_SPAN, SH_YE, SH_QRY, SH_CHK, SH_LOG,
    SH_HELP, SUB_SHEETS, SPARE_NAMES, N_CO, N_ACC, N_LEDGER, TWO_ROWS, K_DAY, ATTR_NO, STALE_WARN, STALE_RED, ACC_COLS,
    BASE_HEAD, BASE_HDR, HOME_SEL, HOME_KPI, YE_SEL, YE_EFF, YE1, YE2, ACC_TYPES, ACC_STATUS, HDR, LEDGER_SHOW,
    RGB_ERR_FILL, RGB_ERR_FONT, RGB_WARN_FONT, PAUSE_TITLE,
    dsel_bad, raw, yr_eff, yr_bad, win_first, win_last, win_before, win, k_aypm, sum_before_year, bref, chain, pause, pause0)
from common import (
    YH, F_HDR, F_TXT, F_TXTB, F_AUTO, F_AUTOB, F_HELP, F_RED, F_KPI_L, F_KPI_V, F_NOTE, F_LINK, F_SEC,
    FILL_KPI, FILL_SUB, FILL_NONE, FILL_AUTO, BD, AC, ACW, AL, ALW, AR, ACC, DATE, INT,
    put, title, header, widths, hide, selector, home_link, link, cf_flags, cf_neg_red, cf_bg, fill, print_setup)

A_IN, A_OUT, A_CAP, A_FLOW, A_TBD, I9 = (ATTR_NO[k] for k in ('收入', '支出', '投资建设', '往来', '待分类', '内部转账'))
DAYS = '0" 天"'
ASHRINK = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)
ACSHRINK = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
F_SUBH = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_BIG = Font(name=YH, sz=12, bold=True, color='FFC00000')   # 再大 make._fit_widths 会把 C 列撑宽
F_SMALL = Font(name=YH, sz=8, color='FF808080')
F_LINK9 = Font(name=YH, sz=9, bold=True, color='FF0563C1', underline='single')
F_STEP = Font(name=YH, sz=10, bold=True, color='FFC00000')
LBL = fill('FFD9E1F2')
HFILL = fill('FF1F3864')
BAND = fill('FFF2F2F2')
WHITE = fill('FFFFFFFF')

# ───────────── 首页版面（行号；model/chk 也用） ─────────────
LAST = 'R'
R_SEL, R_ALERT, R_KH, R_K = 3, 4, 5, (6, 7, 8)              # 第 6/7/8 行＝当天/本月/本年
CO_SEC, CO_HDR, CO_R1 = 10, 11, 12
CO_R2 = CO_R1 + N_CO - 1                                    # 41
CO_TOT = CO_R2 + 1                                          # 42
MIN_CO_VIS = 13                                             # 右边 12 个月表＋合计要 13 行
MON_R1, MON_TOT = CO_R1, CO_R1 + 12                         # 12..23、24
TWO_SEC = CO_TOT + 2                                        # 44
TWO_HDR, TWO_R1 = TWO_SEC + 1, TWO_SEC + 2                  # 45、46
TWO_R2 = TWO_R1 + TWO_ROWS - 1                              # 117
TWO_TOT = TWO_R2 + 1                                        # 118
CHART_AT = f'K{TWO_HDR}'
NAV_SEC = TWO_HDR + 14                                      # 59
CO_COLS = dict(公司='A', 对公='B', 私户='C', 现金='D', 其他='E', 在用小计='F', 暂停户='G', 公司合计='H', 本年经营净额='I')
MON_COLS = dict(月份='K', 经营收入='L', 经营支出='M', 经营净额='N', 投资建设='O', 往来净额='P', 内部净额='Q', 月末资金='R')
TWO_COLS = dict(名称='A', 类型='B', 余额='C', 本月收入='D', 本月支出='E', 最后一笔='F', 几天没记='G', 备注='H')
KPI_ITEMS = [('经营收入', 'F'), ('经营支出', 'G'), ('经营净额', 'H'), ('投资建设', 'I'), ('往来流入', 'K'), ('往来流出', 'L'),
             ('往来净额', 'M'), ('待分类', 'N'), ('内部转入', 'O'), ('内部转出', 'P'), ('内部净额', 'Q'), ('资金净增加', 'R')]
KPI = dict(KPI_ITEMS)
ALERT_SLOTS = ['A', 'B:C', 'D:E', 'F:G', 'H:I', 'K:L', 'M:N', 'O:P', 'Q:R']
ALERT_KEYS = ['问题', '提醒', '没配上', '待分类', '没记账', '按摘要', '跳期', '自检', '新本']
ALERT = {k: s.split(':')[0] + str(R_ALERT) for k, s in zip(ALERT_KEYS, ALERT_SLOTS)}
# 隐藏标量（Z 列值、Y 列名字）
SC_KEYS = ['月初', '月窗首', '月窗末', '年窗首', '年窗末', '余额窗末', '开账月码', '最后月码', '跳期笔数', '没记账户数']
SC = {k: f'Z{i + 1}' for i, k in enumerate(SC_KEYS)}
SCA = {k: f'$Z${i + 1}' for i, k in enumerate(SC_KEYS)}
H1, H2 = 2, N_ACC + 1                                       # 每个账户：第 a+1 行
HC = dict(号='AB', 组='AC', 公司号='AD', 类型号='AE', 状态号='AF', 余额='AG', 月收='AH', 月支='AI', 最后键='AJ', 最后日='AK',
          天数='AL', 原名='AM', 备注='AN', 最晚='AO')
MAPC = dict(键='AP', 原名='AQ')
MAP1, MAP2 = 2, 1 + len(SUB_SHEETS)
MH = dict(月码='AS', 窗首='AT', 窗末='AU', 显示='AV', 线窗末='AW', 线='AX')
TH = dict(键='BA', 号='BB', 组='BC')
HIDE_COLS = ['Y', 'Z'] + list(HC.values()) + list(MAPC.values()) + ['AR'] + list(MH.values()) + ['AY', 'AZ'] + list(TH.values())
NAV_LINKS = [SH_LEDGER, SH_BASE, SH_CHK, SH_LOG, SH_DAY, SH_MON, SH_CO, SH_CLS, SH_ITEMCO, SH_IE, SH_SPAN, SH_YE, SH_QRY, SH_HELP]
# —— v3.1 录入模式、横幅（设计 §15.2.2、§15.1.5）——
ENTRY = 'P_录入模式=1'
ALERT_NEAR = N_LEDGER // 10                                 # 还能录不到容量的 10% → △ 快满（跟台帐 F1 同一条线）
BANNER = 'J1'                                               # 第 1 行右边横幅（合并到 R1；第 1 行 make 不查 pause）
B_ANCHOR = '※ 台帐的表头或右边的 Q 列（勿删）被删了：所有数都没算。马上撤销（Ctrl+Z）；已经存盘了找我们'
B_SHIFT = '※ 台帐被插/删了单元格（不是整行），后面的数全乱了：马上撤销；插删一律点左边行号选整行'
B_FULL = '※ 台帐满了：年底开新本；年中满了找我们放大'
B_ENTRY = '【录入模式】首页只算资金总额（全部流水）和 ※△ 数，子表、报表、核对表先不算；录完去基础资料①清空'
SEC_ENTRY = '录入模式中先不算（基础资料①「录入模式」清空就恢复）'
TBD_REF = f'{LEDGER_SHOW["属性"]}{HDR}'                     # 「待分类」点过去：台帐 M 列（属性）表头，点筛选箭头只勾「待分类」
TBD_TIP = '跳到【资金台帐】M 列（属性）：点表头的筛选箭头，只勾「待分类」，就是这几行；改 C 列收费项，或去基础资料④给收费项定属性'


def banner_formula():
    """首页横幅（J1）：设计 §15.1.5 先后只显示最要紧的一条；只引用 _参 的 P_xx（PAR_SEGCHK 准用）；录入模式也算"""
    c = chain([
        ('P_锚坏=1', f'"{B_ANCHOR}"'),
        ('OR(P_错位>0,P_笔数差<0)', f'"{B_SHIFT}"'),
        ('P_没算>0', '"※ 台帐有 "&P_没算&" 行没算进来（看核对表「台帐分段」）"'),
        ('P_可录<=0', f'"{B_FULL}"'),
        (f'P_可录<{ALERT_NEAR}', '"△ 台帐还能录约 "&P_可录&" 行：年底开新本；年中满了找我们放大"'),
        (ENTRY, f'"{B_ENTRY}"'),
    ], '""')
    return f'=IFERROR({c},"{B_ANCHOR}")'


def entry_or(text, f):
    """留着算的段标题：录入模式时换成一句话（登记 _pause_exempt）"""
    return f'=IF({ENTRY},"{text}",{f[1:] if f.startswith("=") else f})'


def hc(k, a_row):
    return f'{HC[k]}{a_row}'


def hrng(k):
    return f'${HC[k]}${H1}:${HC[k]}${H2}'


def md(x):
    return f'MONTH({x})&"/"&DAY({x})'


def dtxt(x):
    return f'YEAR({x})&"/"&MONTH({x})&"/"&DAY({x})'


def group_name():
    return 'IF(P_集团名称="","集团",P_集团名称)'


def helper(ws, coord, v, fmt=None):
    c = ws[coord]
    c.value = v
    c.font = F_HELP
    if fmt:
        c.number_format = fmt
    return c


def merge_put(ws, rng, v, font, fill_, fmt=None, align=AC, border=True):
    ws.merge_cells(rng)
    m = re.fullmatch(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', rng)
    c = put(ws, f'{m.group(1)}{m.group(2)}', v, font, fill_, fmt=fmt, align=align, border=border)
    if border:
        for ci in range(CI(m.group(1)), CI(m.group(3)) + 1):
            for r in range(int(m.group(2)), int(m.group(4)) + 1):
                ws.cell(row=r, column=ci).border = BD
    return c


def slot(ws, spec, row, v, font, fill_, fmt=None, align=AC, border=True):
    """spec＝'A' 或 'B:C'（同一行合并）"""
    if ':' in spec:
        a, b = spec.split(':')
        return merge_put(ws, f'{a}{row}:{b}{row}', v, font, fill_, fmt=fmt, align=align, border=border)
    return put(ws, f'{spec}{row}', v, font, fill_, fmt=fmt, align=align, border=border)


def type_txt(x):
    return 'CHOOSE(' + x + ',' + ','.join(f'"{t}"' for t in ACC_TYPES) + ')'


def status_txt(x):
    return 'CHOOSE(' + x + ',' + ','.join(f'"{t}"' for t in ACC_STATUS) + ')'


# ═════════════════════════════ 首页 ═════════════════════════════
def _counts(ctx):
    """生成时的行位：公司表显示几行、两级清单用了几行（照 _表 两级候选键的规则从 ctx 推）"""
    cos = [data.key(c.get('公司简称')) for c in ctx['companies'] if data.nm(c.get('公司简称'))]
    n_co = len(ctx['companies'])
    accs = [a for a in ctx['accounts'] if data.nm(a.get('账户名称'))]
    paused = [a for a in accs if data.nm(a.get('状态')) in ACC_STATUS[1:]]            # 跟 _表 状态号一样：暂停/销户＞1，空或别的字＝在用
    active = [a for a in accs if a not in paused]
    co_with = {data.key(a.get('公司')) for a in active if data.key(a.get('公司')) in cos}
    orphan = [a for a in active if data.key(a.get('公司')) not in cos]
    used = len(co_with) + len(accs) + (1 if orphan else 0) + (1 if paused else 0)
    return max(n_co, MIN_CO_VIS), used


def nav_plan(ctx):
    """导航：26 张子表按公司分组（每家公司另起一行，一行 2 个）→ [([(表名, 原表名)…], 组号)]"""
    co_of = {}
    for a in ctx['accounts']:
        if data.nm(a.get('子表名')):
            co_of[data.nm(a.get('子表名'))] = data.nm(a.get('公司'))
    groups = []
    for s, orig, _cap in SUB_SHEETS:
        co = co_of.get(s) or s.split('-')[0]
        if groups and groups[-1][0] == co:
            groups[-1][1].append((s, orig))
        else:
            groups.append((co, [(s, orig)]))
    rows = []
    for gi, (_co, items) in enumerate(groups):
        for i in range(0, len(items), 2):
            rows.append((items[i:i + 2], gi))
    return rows


def build_home(ws, ctx):
    vis_co, used_two = _counts(ctx)
    widths(ws, dict(A=26, **{CL(i): 15 for i in range(2, 10)}, J=1.5, K=13.5, **{CL(i): 15 for i in range(12, 19)}))
    ws.sheet_view.zoomScale = 85
    ws._pause_exempt = []
    tip = ('老板看这一页：所有数都截到 C3「截止日期」（空着＝最后一笔流水那天；基础资料①「管理年度」填了上一年＝那年 12/31），按日期算，跟台帐上下顺序无关。'
           '资金＝各账户的余额，按账户所属公司汇总，暂停/销户的账户单列；经营收入、经营支出只算收费项属性是「收入」「支出」的，'
           '投资建设、往来、待分类、内部转账单列（本年内部净额应该是 0）；各公司的本年经营净额按「归属公司」算。'
           '第 4 行是要处理的事，点一下就跳过去；「待分类」那一格点过去是台帐的属性列（M 列）：点筛选箭头只勾「待分类」就是那几行。'
           '各账户记到哪天看下面「最后一笔」「几天没记」：在用的账户超过 7 天琥珀色、超过 30 天红色。'
           '基础资料①「录入模式」填了「是」时，首页只算资金总额（全部流水）和 ※△ 数，别的先空着，录完清空就恢复。')
    title(ws, '=""', LAST, tip, h2=60)
    # 第 1 行：标题 A1:I1 ＋ 横幅 J1:R1（设计 §15.1.5 先后，录入模式也算；第 1 行不归 pause() 管）
    ws.unmerge_cells(f'A1:{LAST}1')
    ws.merge_cells('A1:I1')
    ws['A1'] = (f'={group_name()}&" 资金看板　·　本本："&P_开账年&" 年（开账 "&{dtxt("P_开账日")}&"）"')
    ws.merge_cells(f'{BANNER}:{LAST}1')
    c = put(ws, BANNER, banner_formula(), Font(name=YH, sz=11, bold=True, color='FFFFFFFF'), HFILL, align=Alignment(
        horizontal='left', vertical='center', wrap_text=True, shrink_to_fit=False), border=False)
    link(c, SH_LEDGER, 'A1')
    rng1 = f'{BANNER}:{LAST}1'
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'LEFT(${BANNER[0]}$1,1)="※"'], stopIfTrue=True,
                                                    font=Font(name=YH, sz=11, bold=True, color=RGB_ERR_FONT), fill=cf_bg(RGB_ERR_FILL)))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'LEN(${BANNER[0]}$1)>0'], stopIfTrue=True,
                                                    font=Font(name=YH, sz=11, bold=True, color=RGB_WARN_FONT), fill=cf_bg('FFFFF2CC')))
    # —— 隐藏标量（录入模式出 0）——
    lym = 'YEAR(P_最后日期)*100+MONTH(P_最后日期)'
    sc = {
        '月初': 'DATE(P_截止年,MONTH(P_截止),1)',
        '月窗首': win_first('P_截止月码'),
        '月窗末': win_last('P_截止月码', SCA['月窗首']),
        '年窗首': win_first('P_截止年*100+1'),
        '年窗末': win_last('P_截止月码', SCA['年窗首']),
        '余额窗末': win_last('P_截止月码'),
        '开账月码': 'YEAR(P_开账日)*100+MONTH(P_开账日)',
        '最后月码': lym,
        '跳期笔数': 'COUNTIF(流_问题,"※ 比前面最晚的一笔*")',
        '没记账户数': f'COUNTIFS({hrng("状态号")},1,{hrng("天数")},">"&{STALE_WARN})',
    }
    for k, f in sc.items():
        helper(ws, SC[k], pause0(f), fmt='yyyy/mm/dd' if k == '月初' else None)
        helper(ws, f'Y{SC[k][1:]}', k)
    # —— 每个账户（第 a+1 行；a＝ROW()−1，不写往下撑的区域）——
    for k, c in HC.items():
        helper(ws, f'{c}1', k)
    bw = SCA['余额窗末']
    wk, wn = win('流_账日键', 1, bw), win('流_净', 1, bw)
    mw = lambda nm: win(nm, SCA['月窗首'], SCA['月窗末'])
    kk = f'(ROW()-{H1 - 1})'
    for r in range(H1, H2 + 1):
        a = f'${HC["号"]}{r}'
        late = f'${HC["最晚"]}{r}'
        helper(ws, hc('号', r), pause0(f'IF(INDEX(账户_有效,{kk})=1,{kk},0)'))
        helper(ws, hc('最晚', r), pause0(f'IF({a}=0,0,MAX(INDEX(块_账最晚,0,{a})))'), fmt='0')
        helper(ws, hc('组', r), pause0(f'IF({a}=0,0,IF(INDEX(账户_状态号,{a})>1,99,IF(INDEX(账户_公司号,{a})>0,INDEX(账户_公司号,{a}),98)))'))
        helper(ws, hc('公司号', r), pause0(f'IF({a}=0,0,INDEX(账户_公司号,{a}))'))
        helper(ws, hc('类型号', r), pause0(f'IF({a}=0,0,INDEX(账户_类型号,{a}))'))
        helper(ws, hc('状态号', r), pause0(f'IF({a}=0,0,INDEX(账户_状态号,{a}))'))
        helper(ws, hc('余额', r), pause0(f'IF({a}=0,0,ROUND(INDEX(账户_期初,{a})+IF({late}<=P_截止,SUM(INDEX(块_净,0,{a})),'
                                         f'SUMIFS({wn},{wk},">="&{a}*{K_DAY},{wk},"<="&({a}*{K_DAY}+P_截止))),2))'))
        for k, val in (('月收', '流_收'), ('月支', '流_支')):
            helper(ws, hc(k, r), pause0(f'IF({a}=0,0,ROUND(SUMIFS({mw(val)},{mw("流_账日键")},">="&({a}*{K_DAY}+{SCA["月初"]}),'
                                        f'{mw("流_账日键")},"<="&({a}*{K_DAY}+P_截止),{mw("流_属性号")},"<>{I9}"),2))'))
        helper(ws, hc('最后键', r), pause0(f'IF({a}=0,0,IF({late}<=P_截止,IF({late}>0,{a}*{K_DAY}+{late},0),'
                                           f'IFERROR(LARGE({wk},COUNTIF({wk},">"&({a}*{K_DAY}+P_截止))+1),0)))'))
        helper(ws, hc('最后日', r), pause0(f'IF({a}=0,0,IF({hc("最后键", r)}>={a}*{K_DAY},{hc("最后键", r)}-{a}*{K_DAY},0))'), fmt='0')
        helper(ws, hc('天数', r), pause(f'IF({hc("最后日", r)}>0,P_截止-{hc("最后日", r)},"")'), fmt='0')
        helper(ws, hc('原名', r), pause(f'IF({a}=0,"",IFERROR(INDEX(${MAPC["原名"]}${MAP1}:${MAPC["原名"]}${MAP2},'
                                       f'MATCH(INDEX(账户_子表名规,{a}),${MAPC["键"]}${MAP1}:${MAPC["键"]}${MAP2},0)),""))'))
        st, sd, on = hc('状态号', r), f'INDEX(账户_停用日,{a})', hc('原名', r)
        helper(ws, hc('备注', r), pause(f'IF({a}=0,"",IF({st}>1,{status_txt(st)}&IF({sd}>0,"（停用 "&{dtxt(sd)}&"）","")&IF({on}<>"","；",""),"")'
                                       f'&IF({on}<>"","原："&{on},""))'))
    # 子表名 → 原表名（生成时写死）
    helper(ws, f'{MAPC["键"]}1', '子表名规')
    helper(ws, f'{MAPC["原名"]}1', '原表名')
    for i, (s, orig, _c) in enumerate(SUB_SHEETS):
        helper(ws, f'{MAPC["键"]}{MAP1 + i}', data.key(s))
        helper(ws, f'{MAPC["原名"]}{MAP1 + i}', orig)

    _top(ws)
    _company(ws, vis_co)
    _months(ws)
    nav_last = _nav(ws, ctx)
    _two(ws, vis_co, max(used_two, nav_last - TWO_R1 + 1))
    hide(ws, *HIDE_COLS)
    for c in HIDE_COLS:
        ws.column_dimensions[c].width = 10
    ws._keep_print = True
    print_setup(ws, None, landscape=True)
    ws.print_area = f'A1:{LAST}{TWO_TOT}'


def _top(ws):
    ex = ws._pause_exempt
    # —— 第 3 行：截止日期（抬头，录入模式照算）——
    sel = f'${HOME_SEL["截止"][0]}${HOME_SEL["截止"][1:]}'
    assert HOME_SEL['截止'] == 'C3'
    selector(ws, 'A3', '截止日期', HOME_SEL['截止'], kind='date',
             prompt='空着＝最后一笔流水那天（基础资料①管理年度填了上一年＝那年 12/31）。改这一格，首页所有数都截到这一天')
    ws.merge_cells('A3:B3')
    note = (f'=IF({dsel_bad(sel)},"※ 看不懂，按 "&{dtxt("P_截止")}&" 算",IF(P_截止<P_开账日,"※ 早于开账日 "&{dtxt("P_开账日")}&"：本本从那天开始",'
            f'"按 "&{dtxt("P_截止")}&" 算"&IF({raw(sel)},"",IF(P_管理年度<YEAR(P_最后日期),"（管理年度 "&P_管理年度&" 年 → 截到年底）",'
            f'"（空着＝最后一笔那天）"))))')
    merge_put(ws, 'D3:F3', note, F_NOTE, FILL_NONE, align=ASHRINK)
    merge_put(ws, 'G3:I3', f'="数据到 "&{dtxt("P_最后日期")}&"（各户记到哪天见下）"', F_SUBH, FILL_NONE, align=ACSHRINK)
    merge_put(ws, f'K3:{LAST}3', '="台帐已用 "&TEXT(P_已用行数,"#,##0")&" 行 / 能录 "&TEXT(P_台帐容量,"#,##0")&" 行（还能录约 "'
                                  '&TEXT(P_可录,"#,##0")&" 行）"', F_SUBH, FILL_NONE, align=ACSHRINK)
    ex += ['D3', 'G3', 'K3']
    cf_flags(ws, 'D3', 'D3')
    ws.row_dimensions[3].height = 24
    # —— 第 4 行：提醒条（录入模式只留 ※/△ 两格：核对_问题、核对_提醒 录入模式照算）——
    n = lambda k: f'N(核对_{k})'
    alerts = {
        '问题': (f'=IF({n("问题")}>0,"※ 台帐有 "&{n("问题")}&" 行要改","✓ 台帐没有 ※")', SH_CHK, 'A1'),
        '提醒': (f'=IF({n("提醒")}>0,"△ "&{n("提醒")}&" 行最好补","✓ 台帐没有 △")', SH_CHK, 'A1'),
        '没配上': (f'=IF({n("没配上")}>0,"△ 内部转账没配上 "&{n("没配上")}&" 笔","✓ 内部转账都配上了")', SH_CHK, 'A1'),
        '待分类': (f'=IF({n("待分类笔数")}>0,"△ 待分类 "&{n("待分类笔数")}&" 笔 净额 "&TEXT({n("待分类金额")},"#,##0")&" 元",'
                f'"✓ 没有待分类")', SH_LEDGER, TBD_REF),
        '没记账': (f'=IF({SCA["没记账户数"]}>0,"△ "&{SCA["没记账户数"]}&" 户超过 {STALE_WARN} 天没记账","✓ 在用的账户 {STALE_WARN} 天内都记了")',
                SH_HOME, f'A{TWO_SEC}'),
        '按摘要': (f'=IF({n("按摘要笔数")}>0,"△ 按摘要认的 "&{n("按摘要笔数")}&" 笔（建议固化）","✓ 没有按摘要认的")', SH_CHK, 'A1'),
        '跳期': (f'=IF({SCA["跳期笔数"]}>0,"※ "&{SCA["跳期笔数"]}&" 笔日期跳期：年份写错了？","✓ 没有日期跳期")', SH_CHK, 'A1'),
        '自检': ('=IF(P_自检>0,"※ 规则自检不过：按摘要认不了","✓ 规则自检正常")', SH_CHK, 'A1'),
        '新本': (f'=IF(YEAR(P_最后日期)>P_开账年,"△ 已经有 "&(P_开账年+1)&" 年的流水：该开新本",IF(P_可录<{ALERT_NEAR},'
               f'"△ 台帐还能录约 "&P_可录&" 行：年底开新本","✓ 台帐容量够用"))', SH_YE, 'A1'),
    }
    keep = ('问题', '提醒')
    for k, spec in zip(ALERT_KEYS, ALERT_SLOTS):
        f, sh, ref = alerts[k]
        c = slot(ws, spec, R_ALERT, f if k in keep else pause(f), Font(name=YH, sz=10, bold=True, color='FF548235'), WHITE, align=ACW)
        link(c, sh, ref)
        if k == '待分类':
            c.hyperlink.tooltip = TBD_TIP
        if k in keep:
            ex.append(ALERT[k])
    rng = f'A{R_ALERT}:{LAST}{R_ALERT}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(A{R_ALERT},1)="※"'], stopIfTrue=True,
                                                   font=Font(name=YH, bold=True, color='FFC00000'), fill=cf_bg('FFFFD7D7')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(A{R_ALERT},1)="△"'], stopIfTrue=True,
                                                   font=Font(name=YH, bold=True, color='FF8B5E00'), fill=cf_bg('FFFFF2CC')))
    ws.row_dimensions[R_ALERT].height = 32
    # —— 第 5～8 行：资金总额卡片（录入模式：C6＝期初合计＋全部流水净额（块表），C7/C8 停）——
    hrow = R_KH
    merge_put(ws, f'A{hrow}:D{hrow}', entry_or('资金总额（录入模式：全部流水，不按截止日）',
                                               f'="资金总额（截止 "&{dtxt("P_截止")}&"，全部账户的余额）"'), F_HDR, HFILL, align=AC)
    assert HOME_KPI['资金总额'] == 'C6'
    merge_put(ws, 'A6:B6', '资金总额', F_KPI_L, FILL_KPI, align=AC)
    merge_put(ws, 'C6:D6', f'=IF({ENTRY},ROUND(P_期初合计+SUM(块_净),2),ROUND(SUM({hrng("余额")}),2))', F_BIG, FILL_KPI, fmt=ACC, align=AR)
    ex += [f'A{hrow}', 'C6']
    merge_put(ws, 'A7:B7', '其中：在用的账户', F_KPI_L, LBL, align=AC)
    merge_put(ws, 'C7:D7', pause(f'ROUND(SUMIF({hrng("状态号")},1,{hrng("余额")}),2)'), F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
    merge_put(ws, 'A8:B8', '其中：暂停 / 销户的账户', F_KPI_L, LBL, align=AC)
    merge_put(ws, 'C8:D8', pause(f'ROUND(SUMIF({hrng("状态号")},">1",{hrng("余额")}),2)'), F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
    cf_neg_red(ws, 'C6:D8', 'C6')
    # —— 当天 / 本月 / 本年 ——
    put(ws, f'E{hrow}', '期间', F_HDR, HFILL, align=ACW)
    heads = {'投资建设': '投资建设\n（花出去）', '待分类': '待分类\n（请定性质）', '内部净额': '内部净额\n（本年应 0）', '资金净增加': '资金净增加\n（全部收－支）'}
    for k, col in KPI_ITEMS:
        put(ws, f'{col}{hrow}', heads.get(k, k), F_HDR, HFILL, align=ACW)
    put(ws, f'J{hrow}', None, F_HDR, HFILL)
    ws.row_dimensions[hrow].height = 32
    lbl = {6: f'="当天 "&{md("P_截止")}', 7: f'="本月 "&MONTH(P_截止)&"/1～"&{md("P_截止")}', 8: f'="本年 1/1～"&{md("P_截止")}'}
    crit = {
        6: (SCA['月窗首'], SCA['月窗末'], lambda w: f'{w("流_月码")},P_截止月码,{w("流_日期")},P_截止'),
        7: (SCA['月窗首'], SCA['月窗末'], lambda w: f'{w("流_月码")},P_截止月码,{w("流_日期")},"<="&P_截止'),
        8: (SCA['年窗首'], SCA['年窗末'], lambda w: f'{w("流_月码")},">="&(P_截止年*100+1),{w("流_月码")},"<="&P_截止月码,{w("流_日期")},"<="&P_截止'),
    }
    for r in R_K:
        put(ws, f'E{r}', pause(lbl[r]), F_SUBH, LBL, align=ACSHRINK)
        put(ws, f'J{r}', None, None, FILL_KPI)
        w1, w2, cf = crit[r]
        w = lambda nm: win(nm, w1, w2)
        cc = cf(w)
        s = lambda val, p=None: f'SUMIFS({w(val)},{cc}' + (f',{w("流_属性号")},{p})' if p is not None else ')')
        vals = {
            '经营收入': f'=ROUND({s("流_净", A_IN)},2)',
            '经营支出': f'=ROUND(-{s("流_净", A_OUT)},2)',
            '经营净额': f'=ROUND({KPI["经营收入"]}{r}-{KPI["经营支出"]}{r},2)',
            '投资建设': f'=ROUND(-{s("流_净", A_CAP)},2)',
            '往来流入': f'=ROUND({s("流_收", A_FLOW)},2)',
            '往来流出': f'=ROUND({s("流_支", A_FLOW)},2)',
            '往来净额': f'=ROUND({KPI["往来流入"]}{r}-{KPI["往来流出"]}{r},2)',
            '待分类': f'=ROUND({s("流_净", A_TBD)},2)',
            '内部转入': f'=ROUND({s("流_收", I9)},2)',
            '内部转出': f'=ROUND({s("流_支", I9)},2)',
            '内部净额': f'=ROUND({KPI["内部转入"]}{r}-{KPI["内部转出"]}{r},2)',
            '资金净增加': f'=ROUND({s("流_净")},2)',
        }
        for k, col in KPI_ITEMS:
            put(ws, f'{col}{r}', pause(vals[k]), F_AUTOB if k in ('经营净额', '资金净增加') else F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        ws.row_dimensions[r].height = 22
    for k in ('经营净额', '往来净额', '待分类', '资金净增加'):
        cf_neg_red(ws, f'{KPI[k]}6:{KPI[k]}8', f'{KPI[k]}6')
    q8 = f'{KPI["内部净额"]}8'
    ws.conditional_formatting.add(q8, FormulaRule(formula=[f'ROUND(N({q8}),2)<>0'], font=Font(name=YH, bold=True, color='FFC00000'),
                                                  fill=cf_bg('FFFFD7D7')))


def _company(ws, vis_co):
    C = CO_COLS
    co_hidden = f'COUNTIF(INDEX(公司_有效,{vis_co + 1}):INDEX(公司_有效,{N_CO}),1)' if vis_co < N_CO else '0'
    orphan = f'COUNTIF({hrng("组")},98)'
    t = (f'="各公司资金（截止 "&{md("P_截止")}&"；资金按账户所属公司，本年经营净额按归属公司）"'
         f'&IF({orphan}>0,"　△ "&{orphan}&" 个账户的公司没认出：合计里有、公司行里没有（去基础资料③改「公司」）","")'
         f'&IF({co_hidden}>0,"　△ 下面有藏起来的公司行（第 {CO_R1 + vis_co}～{CO_R2} 行）：选中这几行，右键「取消隐藏」","")')
    merge_put(ws, f'A{CO_SEC}:I{CO_SEC}', entry_or(f'各公司资金：{SEC_ENTRY}', t), F_SEC, HFILL, align=ASHRINK)
    ws._pause_exempt.append(f'A{CO_SEC}')
    ws.row_dimensions[CO_SEC].height = 22
    hdr = dict(公司='公司', 对公='对公账户', 私户='私户', 现金='现金', 其他='其他', 在用小计='在用小计', 暂停户='暂停 / 销户',
               公司合计='公司合计', 本年经营净额='本年经营净额\n（按归属公司）')
    header(ws, CO_HDR, [(C[k], hdr[k]) for k in C], '首页', height=32)
    wy = lambda nm: win(nm, SCA['年窗首'], SCA['年窗末'])
    Y = 'P_截止年'
    k = f'(ROW()-{CO_R1 - 1})'                                 # 第几家公司（不写往下撑的区域）
    for r in range(CO_R1, CO_R2 + 1):
        nm = f'${C["公司"]}{r}'
        put(ws, f'{C["公司"]}{r}', pause(f'IF(INDEX(公司_有效,{k})=1,INDEX(公司_名称,{k}),"")'), F_TXTB, FILL_NONE, align=ASHRINK)
        for t_no, key in enumerate(('对公', '私户', '现金', '其他'), 1):
            put(ws, f'{C[key]}{r}', pause(f'IF({nm}="","",SUMIFS({hrng("余额")},{hrng("公司号")},{k},{hrng("类型号")},{t_no},{hrng("状态号")},1))'),
                F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'{C["在用小计"]}{r}', pause(f'IF({nm}="","",ROUND(SUM({C["对公"]}{r}:{C["其他"]}{r}),2))'), F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'{C["暂停户"]}{r}', pause(f'IF({nm}="","",SUMIFS({hrng("余额")},{hrng("公司号")},{k},{hrng("状态号")},">1"))'),
            F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'{C["公司合计"]}{r}', pause(f'IF({nm}="","",ROUND({C["在用小计"]}{r}+{C["暂停户"]}{r},2))'), F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'{C["本年经营净额"]}{r}', pause(f'IF({nm}="","",ROUND(SUMIFS({wy("流_净")},{wy("流_归年属月键")},">="&{k_aypm(k, Y, A_IN, 1)},'
                                              f'{wy("流_归年属月键")},"<="&{k_aypm(k, Y, A_OUT, 12)},{wy("流_日期")},"<="&P_截止),2))'),
            F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
        ws.row_dimensions[r].height = 18
        if r >= CO_R1 + vis_co:
            ws.row_dimensions[r].hidden = True
    rt = CO_TOT
    put(ws, f'{C["公司"]}{rt}', '合计（全部账户）', F_RED, FILL_KPI, align=AC)
    for t_no, key in enumerate(('对公', '私户', '现金', '其他'), 1):
        put(ws, f'{C[key]}{rt}', pause(f'ROUND(SUMIFS({hrng("余额")},{hrng("类型号")},{t_no},{hrng("状态号")},1),2)'), F_RED, FILL_KPI, fmt=ACC,
            align=AR)
    put(ws, f'{C["在用小计"]}{rt}', pause(f'ROUND(SUM({C["对公"]}{rt}:{C["其他"]}{rt}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{C["暂停户"]}{rt}', pause(f'ROUND(SUMIF({hrng("状态号")},">1",{hrng("余额")}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{C["公司合计"]}{rt}', pause(f'ROUND({C["在用小计"]}{rt}+{C["暂停户"]}{rt},2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{C["本年经营净额"]}{rt}', pause(f'${KPI["经营净额"]}$8'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    ws.row_dimensions[rt].height = 20
    for k_ in ('在用小计', '公司合计', '本年经营净额', '对公', '私户', '现金', '其他', '暂停户'):
        cf_neg_red(ws, f'{C[k_]}{CO_R1}:{C[k_]}{rt}', f'{C[k_]}{CO_R1}')


def _months(ws):
    M = MON_COLS
    merge_put(ws, f'K{CO_SEC}:{LAST}{CO_SEC}', entry_or(f'每月：{SEC_ENTRY}', '=P_截止年&" 年每月（经营口径；到 "&' + md('P_截止') +
                                                         '&" 为止；月末资金＝全部账户的余额）"'), F_SEC, HFILL, align=ASHRINK)
    ws._pause_exempt.append(f'K{CO_SEC}')
    hdr = dict(月份='月份', 经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设', 往来净额='往来净额',
               内部净额='内部净额', 月末资金='月末资金总额')
    header(ws, CO_HDR, [(M[k], hdr[k]) for k in M], '首页', height=32)
    for k, c in MH.items():
        helper(ws, f'{c}{CO_HDR}', k)
    Y = 'P_截止年'
    for m in range(1, 13):
        r = MON_R1 + m - 1
        ym = f'${MH["月码"]}{r}'
        helper(ws, f'{MH["月码"]}{r}', pause0(f'{Y}*100+ROW()-{MON_R1 - 1}'))
        helper(ws, f'{MH["窗首"]}{r}', pause0(win_first(ym)))
        helper(ws, f'{MH["窗末"]}{r}', pause0(win_last(ym, f'{MH["窗首"]}{r}')))
        w = lambda nm: win(nm, f'${MH["窗首"]}{r}', f'${MH["窗末"]}{r}')
        helper(ws, f'{MH["显示"]}{r}', pause0(f'IF(AND({ym}>={SCA["开账月码"]},{ym}<=P_截止月码,OR({ym}<={SCA["最后月码"]},'
                                              f'COUNTIF({w("流_月码")},{ym})>0)),1,0)'))
        lym = f'MIN({ym},P_截止月码)'
        helper(ws, f'{MH["线窗末"]}{r}', pause0(win_last(lym)))
        wl = lambda nm: win(nm, 1, f'${MH["线窗末"]}{r}')
        if m == 1:                                          # 1 月：期初合计＋1 月底（不过截止）以前的全部
            line = f'ROUND(P_期初合计+SUMIFS({wl("流_净")},{wl("流_月码")},"<="&{lym},{wl("流_日期")},"<="&P_截止),2)'
        else:                                               # 2～12 月：上个月＋本月（只扫本月窗口；截止以后的月份＝截止那个月）
            line = (f'ROUND(${MH["线"]}{r - 1}+IF({ym}<=P_截止月码,SUMIFS({w("流_净")},{w("流_月码")},{ym},{w("流_日期")},"<="&P_截止),0),2)')
        helper(ws, f'{MH["线"]}{r}', pause(line))
        sh = f'${MH["显示"]}{r}=0'
        cc = f'{w("流_月码")},{ym},{w("流_日期")},"<="&P_截止'
        s = lambda val, p: f'SUMIFS({w(val)},{cc},{w("流_属性号")},{p})'
        vals = dict(经营收入=f'ROUND({s("流_净", A_IN)},2)', 经营支出=f'ROUND(-{s("流_净", A_OUT)},2)',
                    经营净额=f'ROUND({M["经营收入"]}{r}-{M["经营支出"]}{r},2)', 投资建设=f'ROUND(-{s("流_净", A_CAP)},2)',
                    往来净额=f'ROUND({s("流_净", A_FLOW)},2)', 内部净额=f'ROUND({s("流_净", I9)},2)', 月末资金=f'${MH["线"]}{r}')
        put(ws, f'{M["月份"]}{r}', f'{m}月', F_TXTB, FILL_NONE, align=AC)
        for k, v in vals.items():
            put(ws, f'{M[k]}{r}', pause(f'IF({sh},"",{v})'), F_AUTOB if k in ('经营净额', '月末资金') else F_AUTO, FILL_NONE, fmt=ACC, align=AR)
    rt = MON_TOT
    put(ws, f'{M["月份"]}{rt}', '合计', F_RED, FILL_KPI, align=AC)
    for k in ('经营收入', '经营支出', '经营净额', '投资建设', '往来净额', '内部净额'):
        put(ws, f'{M[k]}{rt}', pause(f'ROUND(SUM({M[k]}{MON_R1}:{M[k]}{rt - 1}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{M["月末资金"]}{rt}', pause('$C$6'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    for k in ('经营净额', '往来净额', '月末资金'):
        cf_neg_red(ws, f'{M[k]}{MON_R1}:{M[k]}{rt}', f'{M[k]}{MON_R1}')
    qn = f'{M["内部净额"]}{rt}'
    ws.conditional_formatting.add(qn, FormulaRule(formula=[f'ROUND(N({qn}),2)<>0'], font=Font(name=YH, bold=True, color='FFC00000'),
                                                 fill=cf_bg('FFFFD7D7')))
    # —— 走势图（Excel 2010 能看：柱＝经营收入/支出，线＝月末资金总额（右轴，引用隐藏的「线」列：未来月份＝截止日余额）——
    merge_put(ws, f'K{TWO_SEC}:{LAST}{TWO_SEC}', entry_or(f'走势图：{SEC_ENTRY}', '=P_截止年&" 年走势：柱＝经营收入、经营支出（左轴）；'
                                                                                  '线＝月末资金总额（右轴）"'), F_SEC, HFILL, align=ASHRINK)
    ws._pause_exempt.append(f'K{TWO_SEC}')
    bar = BarChart()
    bar.type, bar.grouping = 'col', 'clustered'
    bar.add_data(Reference(ws, min_col=CI(M['经营收入']), max_col=CI(M['经营支出']), min_row=CO_HDR, max_row=MON_R1 + 11),
                 titles_from_data=True)
    bar.set_categories(Reference(ws, min_col=CI(M['月份']), min_row=MON_R1, max_row=MON_R1 + 11))
    bar.y_axis.title = '经营收支'
    bar.y_axis.number_format = '#,##0'
    bar.y_axis.majorGridlines = None
    bar.x_axis.delete = False
    bar.y_axis.delete = False
    bar.gapWidth = 60
    line = LineChart()
    line.add_data(Reference(ws, min_col=CI(MH['线']), min_row=MON_R1, max_row=MON_R1 + 11), titles_from_data=False)
    line.series[0].tx = SeriesLabel(v='月末资金总额')
    line.series[0].smooth = False
    line.y_axis.axId = 200
    line.y_axis.title = '月末资金'
    line.y_axis.number_format = '#,##0'
    line.y_axis.crosses = 'max'
    line.y_axis.delete = False
    line.y_axis.majorGridlines = None
    cat = f"'{SH_HOME}'!${M['月份']}${MON_R1}:${M['月份']}${MON_R1 + 11}"
    for sr in list(bar.series) + list(line.series):           # 月份是文字（1月…12月）：分类用 strRef（numRef 指向文字 Excel 会显示成 1..12）
        sr.cat = AxDataSource(strRef=StrRef(f=cat))
    bar += line
    bar.visible_cells_only = False                          # 「线」在隐藏列里，要画隐藏格
    bar.legend.position = 'b'
    bar.width, bar.height = 23.0, 7.8
    ws.add_chart(bar, CHART_AT)
    ws._home_chart = bar


def _nav(ws, ctx):
    r = NAV_SEC
    merge_put(ws, f'K{r}:{LAST}{r}', '导航（点名字跳过去；每张表第 3 行都有「← 回首页」）', F_SEC, HFILL, align=AL)
    slots = ['K:L', 'M:N', 'O:P', 'Q:R']
    r += 1
    for i in range(0, len(NAV_LINKS), 4):
        for spec, sh in zip(slots, NAV_LINKS[i:i + 4]):
            c = slot(ws, spec, r, sh, F_LINK, WHITE, align=ACSHRINK)
            link(c, sh, 'A1')
        r += 1
    merge_put(ws, f'K{r}:{LAST}{r}', '账户子表（按公司排；小字是原来日记账的表名）', F_SUBH, LBL, align=AL)
    r += 1
    for items, gi in nav_plan(ctx):
        band = BAND if gi % 2 == 0 else WHITE
        for (s, orig), (c1, c2, c3) in zip(items, (('K', 'M', 'N'), ('O', 'Q', 'R'))):
            c = merge_put(ws, f'{c1}{r}:{c2}{r}', s, F_LINK9, band, align=ASHRINK)
            link(c, s, 'A1')
            put(ws, f'{c3}{r}', f'原：{orig}', F_SMALL, band, align=ASHRINK)
        if len(items) == 1:
            merge_put(ws, f'O{r}:Q{r}', None, F_SMALL, band)
            put(ws, f'R{r}', None, F_SMALL, band)
        r += 1
    put(ws, f'K{r}', '备用子表', F_SUBH, LBL, align=ACSHRINK)
    for spec, sh in zip(['L', 'M:N', 'O:P', 'Q:R'], SPARE_NAMES):
        c = slot(ws, spec, r, sh, F_LINK9, WHITE, align=ACSHRINK)
        link(c, sh, 'A1')
    # 复核 F4：别教「把标签改成账户名」（首页、核对表的链接按表名跳，改名就断）；用的是哪个账户看备用表 A1 标题
    merge_put(ws, f'K{r + 1}:{LAST}{r + 1}', '备用子表给以后新开的卡用：在它的 D3 选账户；表名别改（首页、核对表的链接按表名跳），'
                                           '用的是哪个账户看它的标题', F_NOTE, FILL_NONE, align=ASHRINK, border=False)
    return r + 1


def _two(ws, vis_co, vis_two):
    T = TWO_COLS
    two_hidden = f'COUNT(账两级_候选键)>{vis_two}'
    t = (f'="各账户余额（截止 "&{md("P_截止")}&"；公司 → 账户，暂停 / 销户的放最后；本月收入、支出不含内部转账）"'
         f'&IF({two_hidden},"　△ 账户多了，下面有藏起来的行：选中第 {TWO_R1 + vis_two}～{TWO_R2} 行，右键「取消隐藏」","")')
    merge_put(ws, f'A{TWO_SEC}:I{TWO_SEC}', entry_or(f'各账户余额：{SEC_ENTRY}', t), F_SEC, HFILL, align=ASHRINK)
    ws._pause_exempt.append(f'A{TWO_SEC}')
    ws.row_dimensions[TWO_SEC].height = 22
    hdr = dict(名称='公司 / 账户', 类型='类型', 余额='余额', 本月收入='本月收入\n（不含内部）', 本月支出='本月支出\n（不含内部）',
               最后一笔='最后一笔\n（截止日以前）', 几天没记='几天没记', 备注='备注')
    header(ws, TWO_HDR, [(T[k], hdr[k]) for k in T] + [('I', None)], '首页', height=32)
    ws.merge_cells(f'H{TWO_HDR}:I{TWO_HDR}')
    for k, c in TH.items():
        helper(ws, f'{c}{TWO_HDR}', k)
    for r in range(TWO_R1, TWO_R2 + 1):
        key, a, c = f'${TH["键"]}{r}', f'${TH["号"]}{r}', f'${TH["组"]}{r}'
        helper(ws, f'{TH["键"]}{r}', pause0(f'IFERROR(SMALL(账两级_候选键,ROW()-{TWO_R1 - 1}),0)'))
        helper(ws, f'{TH["号"]}{r}', pause0(f'MOD({TH["键"]}{r},100)'))
        helper(ws, f'{TH["组"]}{r}', pause0(f'INT({TH["键"]}{r}/100)'))
        nm = (f'=IF({key}=0,"",IF({a}>0,"　"&INDEX(账户_名称,{a}),IF({c}=98,"公司没认出的账户（去基础资料③改「公司」）",'
              f'IF({c}=99,"暂停 / 销户的账户",INDEX(公司_名称,{c})))))')
        put(ws, f'{T["名称"]}{r}', pause(nm), F_TXT, FILL_NONE, align=ASHRINK)
        tt = type_txt(f'INDEX({hrng("类型号")},{a})')
        put(ws, f'{T["类型"]}{r}', pause(f'IF({key}=0,"",IF({a}>0,IFERROR({tt},""),"小计"))'), F_TXT, FILL_NONE, align=AC)
        for col, hk in (('余额', '余额'), ('本月收入', '月收'), ('本月支出', '月支')):
            put(ws, f'{T[col]}{r}', pause(f'IF({key}=0,"",IF({a}>0,INDEX({hrng(hk)},{a}),SUMIF({hrng("组")},{c},{hrng(hk)})))'),
                F_AUTOB if col == '余额' else F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'{T["最后一笔"]}{r}', pause(f'IF({a}=0,"",IF(INDEX({hrng("最后日")},{a})>0,INDEX({hrng("最后日")},{a}),"没有流水"))'),
            F_AUTO, FILL_NONE, fmt=DATE, align=AC)
        put(ws, f'{T["几天没记"]}{r}', pause(f'IF({a}=0,"",INDEX({hrng("天数")},{a}))'), F_AUTO, FILL_NONE, fmt=DAYS, align=AC)
        merge_put(ws, f'H{r}:I{r}', pause(f'IF({a}=0,"",INDEX({hrng("备注")},{a}))'), F_NOTE, FILL_NONE, align=ASHRINK)
        ws.row_dimensions[r].height = 18
        if r >= TWO_R1 + vis_two:
            ws.row_dimensions[r].hidden = True
    rt = TWO_TOT
    put(ws, f'A{rt}', '合计（全部账户）', F_RED, FILL_KPI, align=AC)
    put(ws, f'B{rt}', None, F_RED, FILL_KPI)
    for col, hk in (('余额', '余额'), ('本月收入', '月收'), ('本月支出', '月支')):
        put(ws, f'{T[col]}{rt}', pause(f'ROUND(SUM({hrng(hk)}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    for col in ('F', 'G'):
        put(ws, f'{col}{rt}', None, F_RED, FILL_KPI)
    merge_put(ws, f'H{rt}:I{rt}', None, F_RED, FILL_KPI)
    rng = f'A{TWO_R1}:I{TWO_R2}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(${TH["键"]}{TWO_R1}<>0,${TH["号"]}{TWO_R1}=0)'],
                                                   font=Font(name=YH, bold=True, color='FF1F3864'), fill=cf_bg('FFDDEBF7')))
    gr = f'{T["最后一笔"]}{TWO_R1}:{T["几天没记"]}{TWO_R2}'
    gg = f'${T["几天没记"]}{TWO_R1}'
    ok = f'${TH["号"]}{TWO_R1}>0,${TH["组"]}{TWO_R1}<>99,ISNUMBER({gg})'
    ws.conditional_formatting.add(gr, FormulaRule(formula=[f'AND({ok},{gg}>{STALE_RED})'], stopIfTrue=True,
                                                  font=Font(name=YH, bold=True, color='FFC00000'), fill=cf_bg('FFFFD7D7')))
    ws.conditional_formatting.add(gr, FormulaRule(formula=[f'AND({ok},{gg}>{STALE_WARN})'], stopIfTrue=True,
                                                  font=Font(name=YH, bold=True, color='FF8B5E00'), fill=cf_bg('FFFFF2CC')))
    cf_neg_red(ws, f'{T["余额"]}{TWO_R1}:{T["余额"]}{rt}', f'{T["余额"]}{TWO_R1}')


# ═════════════════════════════ 年末余额 ═════════════════════════════
YC = YE1['列']
YP = dict(年初窗末='O1', 年末窗末='O2', 号='O', 最晚='P')
PASTE_FIELDS = [f for f in ACC_COLS if f not in ('序号', '核对')]          # 账户名称～备注（B..M）
assert [ACC_COLS[f] for f in PASTE_FIELDS] == [CL(i) for i in range(2, 14)]


def build_ye(ws, ctx):
    Y = '$' + YE_EFF['年份'][0] + '$' + YE_EFF['年份'][1:]
    widths(ws, dict(A=6, B=30, C=24, D=14, E=14, F=24, G=15, H=15, I=15, J=34, K=13, L=26, M=24))
    tip = ('每个账户某一年的年末余额（按日期算：开账期初＋这一年 12/31 以前的收－支）。年份空着＝管理年度。'
           '下面的粘贴区是开新本用的：跟【基础资料】③账户一列对一列（B..M），期初余额已经换成年末余额，照第 50 行的五步做。'
           '年中新开的账户在新本里期初就是今年的年末余额，不用再录「账户转入」。全部自动算，不用填。'
           '基础资料①「录入模式」填了「是」时本表先不算（全空），清空就恢复。')
    title(ws, '=""', 'L', tip, h2=46)
    # 第 1 行标题：录入模式时换成 PAUSE_TITLE（红字；第 1 行不归 pause() 管）
    ws['A1'] = f'=IF({ENTRY},"{PAUSE_TITLE}",{group_name()}&" "&{Y}&" 年末各账户余额")'
    ws.conditional_formatting.add('A1', FormulaRule(formula=[ENTRY], font=Font(name=YH, sz=15, bold=True, color=RGB_ERR_FONT),
                                                    fill=cf_bg(RGB_ERR_FILL)))
    # M1＝生效年份（年末_年份；第 1 行，录入模式也照算：只看 B3 和 P_管理年度）
    put(ws, YE_EFF['年份'], '=' + yr_eff(f'${YE_SEL["年份"][0]}${YE_SEL["年份"][1:]}'), Font(name=YH, sz=8, color='FF1F3864'),
        fill('FF1F3864'), border=False)
    ws[YE_EFF['年份']].number_format = ';;;'
    sel = f'${YE_SEL["年份"][0]}${YE_SEL["年份"][1:]}'
    selector(ws, 'A3', '年份', YE_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    merge_put(ws, 'C3:D3', pause(f'IF({yr_bad(sel)},"※ 看不懂，按 "&{Y}&" 年算","按 "&{Y}&" 年算"&IF({raw(sel)},"","（管理年度）"))'), F_NOTE,
              FILL_NONE, align=ASHRINK)
    st = (f'IF({Y}<P_开账年,"※ 本本从 "&P_开账年&" 年开始："&{Y}&" 年末＝开账期初，更早的看上一本",IF({Y}>YEAR(P_最后日期),'
          f'"△ "&{Y}&" 年还没有流水：年末余额就是现在的余额",IF(P_最后日期<DATE({Y},12,31),"△ "&{Y}&" 年还没过完：年末余额是到 "&'
          f'{dtxt("P_最后日期")}&" 为止的，记完 12/31 再开新本","✓ "&{Y}&" 年已经记到 12/31，可以开新本")))')
    merge_put(ws, 'E3:J3', pause(st), F_TXTB, FILL_NONE, align=ASHRINK)
    cf_flags(ws, 'C3:J3', 'C3')
    home_link(ws, 'L3')
    ws.row_dimensions[3].height = 24
    helper(ws, YP['年初窗末'], pause0(win_before(f'{Y}*100+1')))
    helper(ws, YP['年末窗末'], pause0(win_last(f'{Y}*100+12')))
    w0 = lambda nm: win(nm, 1, f'${YP["年初窗末"][0]}${YP["年初窗末"][1:]}')
    w1 = lambda nm: win(nm, 1, f'${YP["年末窗末"][0]}${YP["年末窗末"][1:]}')
    hdr = dict(序号='序号', 账户名称='账户名称', 公司='公司', 类型='类型', 状态='状态', 开账期初='开账期初\n（基础资料③）',
               年初余额='年初余额\n（上年末）', 本年净增加='本年净增加', 年末余额='年末余额\n（12/31）', 备注='备注')
    header(ws, YE1['表头行'], [(YC[k], hdr[k]) for k in YC], '报表', height=34)
    helper(ws, f'{YP["号"]}{YE1["表头行"]}', '账户号')
    helper(ws, f'{YP["最晚"]}{YE1["表头行"]}', '最晚')
    r1, rt = YE1['首行'], YE1['合计行']
    kk = f'(ROW()-{r1 - 1})'                                  # 第几个账户（不写往下撑的区域；本表有保护，不插行）
    for r in range(r1, rt):
        a = f'${YP["号"]}{r}'
        helper(ws, f'{YP["号"]}{r}', pause0(f'IF(INDEX(账户_有效,{kk})=1,{kk},0)'))
        late = f'${YP["最晚"]}{r}'
        helper(ws, late[1:], pause0(f'IF({a}=0,0,MAX(INDEX(块_账最晚,0,{a})))'), fmt='0')
        c = lambda k: f'{YC[k]}{r}'
        stn, sd = f'INDEX(账户_状态号,{a})', f'INDEX(账户_停用日,{a})'
        put(ws, c('序号'), pause(f'IF({a}=0,"",COUNTIF(INDEX(账户_有效,1):INDEX(账户_有效,{kk}),1))'), F_AUTO, FILL_NONE, fmt=INT, align=AC)
        put(ws, c('账户名称'), pause(f'IF({a}=0,"",INDEX(账户_名称,{a}))'), F_TXT, FILL_NONE, align=ASHRINK)
        put(ws, c('公司'), pause(f'IF({a}=0,"",IF(INDEX(账户_公司号,{a})>0,INDEX(公司_名称,INDEX(账户_公司号,{a})),"（公司没认出）"))'), F_TXT,
            FILL_NONE, align=ACSHRINK)
        put(ws, c('类型'), pause(f'IF({a}=0,"",IFERROR({type_txt(f"INDEX(账户_类型号,{a})")},""))'), F_TXT, FILL_NONE, align=AC)
        put(ws, c('状态'), pause(f'IF({a}=0,"",IFERROR({status_txt(stn)},"")&IF(AND({stn}>1,{sd}>0),"（"&{dtxt(sd)}&"）",""))'), F_TXT,
            FILL_NONE, align=ACSHRINK)
        put(ws, c('开账期初'), pause(f'IF({a}=0,"",INDEX(账户_期初,{a}))'), F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, c('年初余额'), pause(f'IF({a}=0,"",ROUND(INDEX(账户_期初,{a})+IF({Y}<=P_开账年,0,IF({late}<DATE({Y},1,1),SUM(INDEX(块_净,0,{a})),'
                                    f'{sum_before_year(w0("流_净"), w0("流_账年属月键"), a, Y)})),2))'),
            F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, c('年末余额'), pause(f'IF({a}=0,"",ROUND(INDEX(账户_期初,{a})+IF({late}<=DATE({Y},12,31),SUM(INDEX(块_净,0,{a})),'
                                    f'{sum_before_year(w1("流_净"), w1("流_账年属月键"), a, f"{Y}+1")}),2))'),
            F_AUTOB, FILL_NONE, fmt=ACC, align=AR)
        put(ws, c('本年净增加'), pause(f'IF({a}=0,"",ROUND({c("年末余额")}-{c("年初余额")},2))'), F_AUTO, FILL_NONE, fmt=ACC, align=AR)
        put(ws, c('备注'), pause(f'IF({a}=0,"",IF({c("年末余额")}<0,"△ 年末余额是负数：查漏录的收入、多录的支出",'
                               f'IF(AND({stn}=3,ROUND({c("年末余额")},2)<>0),"△ 已销户但余额不是 0：查一下销户转出那笔",'
                               f'IF(AND({stn}=1,INDEX(账户_公司号,{a})=0),"△ 公司没认出：去基础资料③改「公司」",""))))'),
            F_NOTE, FILL_NONE, align=ASHRINK)
    put(ws, f'{YC["序号"]}{rt}', None, F_RED, FILL_KPI)
    merge_put(ws, f'{YC["账户名称"]}{rt}:{YC["状态"]}{rt}', '合计', F_RED, FILL_KPI)
    for k in ('开账期初', '年初余额', '本年净增加', '年末余额'):
        put(ws, f'{YC[k]}{rt}', pause(f'ROUND(SUM({YC[k]}{r1}:{YC[k]}{rt - 1}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{YC["备注"]}{rt}', None, F_RED, FILL_KPI)
    cf_flags(ws, f'{YC["备注"]}{r1}:{YC["备注"]}{rt}', f'{YC["备注"]}{r1}')
    for k in ('年初余额', '年末余额', '本年净增加'):
        cf_neg_red(ws, f'{YC[k]}{r1}:{YC[k]}{rt}', f'{YC[k]}{r1}')
    _paste(ws, Y)
    hide(ws, 'O', 'P')


def _paste(ws, Y):
    t, s, h, p1 = YE2['标题行'], YE2['步骤行'], YE2['表头行'], YE2['首行']
    p2 = p1 + N_ACC - 1
    b1 = BASE_HDR['账户'] + 1
    merge_put(ws, f'A{t}:M{t}', pause(f'"复制到新本【基础资料】③账户：选中 B{p1}:M{p2} → 复制 → 新本基础资料 B{b1} → 选择性粘贴「数值」（"&{Y}&'
                                      f'" 年末余额 → "&({Y}+1)&" 年期初）"'), F_SEC, fill('FF30561E'), align=AL)
    steps = (f'"开新本五步（顺序别错）：① 本本另存为新文件（如「…"&({Y}+1)&".xlsx」），下面都在新文件里做；② 这张表 B3 填 "&{Y}&" → 选中 B{p1}:M{p2} → 复制 → '
             f'【基础资料】③账户 B{b1}（第一个账户的名称格）→ 选择性粘贴「数值」；③【基础资料】①开账日改 "&({Y}+1)&"-01-01、管理年度清空；'
             f'④【资金台帐】点日期列的筛选箭头只选 "&{Y}&" 年 → 选中这些行的 A:K 和 P 列按 Delete 清内容（别删行，别动 H、L:O 公式列），'
             f'再按日期升序排一次；⑤【数据整理记录】清空，或写「"&{Y}&" 年的导入记录见上一本」。上一年的报表看上一本。"')
    merge_put(ws, f'A{s}:M{s}', pause(steps), F_STEP, fill('FFFFF7E6'), align=ALW)
    ws.row_dimensions[s].height = 62
    heads = BASE_HEAD['账户']
    header(ws, h, [('A', '第几行')] + [(ACC_COLS[f], heads[f]) for f in PASTE_FIELDS], '报表', height=46)
    r1 = YE1['首行']
    for i in range(N_ACC):
        r = p1 + i
        tr = r1 + i
        a = f'${YP["号"]}{tr}'
        put(ws, f'A{r}', i + 1, F_NOTE, FILL_NONE, align=AC)
        for f in PASTE_FIELDS:
            col = ACC_COLS[f]
            if f == '账户名称':
                v = f'IF({a}=0,"",INDEX(账户_名称,{a}))'
            elif f == '期初余额':
                v = f'IF({a}=0,"",ROUND(N({YC["年末余额"]}{tr}),2))'
            else:
                x = f'INDEX({bref("账户", f)},{a}+1)'             # 基础资料③原文：名称 基_账户_字段（不直接写「基础资料!」）
                v = f'IF({a}=0,"",IF({x}&""="","",{x}))'
            fmt = ACC if f == '期初余额' else ('yyyy/mm/dd' if f == '停用日' else None)
            put(ws, f'{col}{r}', pause(v), F_AUTOB if f in ('账户名称', '期初余额') else F_AUTO, FILL_NONE, fmt=fmt,
                align=AR if f == '期初余额' else ASHRINK)


def build(wb, ctx):
    build_home(wb[SH_HOME], ctx)
    build_ye(wb[SH_YE], ctx)
