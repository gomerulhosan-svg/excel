# -*- coding: utf-8 -*-
"""A074 集团多账户资金台帐 · 账户子表（26 张）、备用子表（2 张）、账户查询（设计.md §6；v3.1 提速改动见 §15.1.7、§15.2、§15.4）。

版面（layout：SUB_* / QRY_*）：
  第 1 行标题（账户名、年份、原表名）；第 2 行说明；第 3 行 回首页｜账户（黄 D3）｜信息｜年份（黄 H3）｜状态（账户查询多 月份 J3）；
  第 4 行 KPI：年初余额 D4｜本年收入 F4｜本年支出 H4（不含内部转账）｜年末/当前余额 J4｜笔数 K4；
  第 5 行表头（A..K 明细；M..S 月份表）；第 6 行「年初余额」；第 7 行起逐笔（像原日记账：月、日、收费项、收付人、摘要、收入、支出、余额…）。
逐笔：隐藏 Z 列＝这一笔在 _流 的第几个位置（链式：每行只 MATCH 一次，从上一笔往后找；整张表只扫一遍 _流；段尾空位键是 0，自动跳过），
  其余各列 INDEX(流_xx, Z)；摘要、备注、对方账户、归属公司、没认出的收费项/收付人原文用 layout.lref_z（整列名称 台帐_字段；
  v3.1：_流 位置 ≠ 台帐第几条，台帐行号用 段_区首＋取数号−1）。Z=0 的行每一列都 IF 挡成空（INDEX(…,0) 会取整列/表头）。
  不用 FILTER/OFFSET/INDIRECT；公式里不直接写「资金台帐!」、不写往下撑的区域（WPS 存盘会一格一条拆开写；序号写 ROW()−6）。
月份表（M5:S19）：按 _流 的键算（账户口径：收入/支出不含内部转账，内部转入/转出单列；月末余额＝上月＋收−支＋入−出），每个月只扫那个月的窗口。
  当年最后日期以后的月份显示空；第 19 行全年（R19＝R6＋N19−O19＋P19−Q19，不取 R18）。
录入模式（v3.1，基础资料①「录入模式」＝是）：链头 Z7 用 pause0 出 0 → 整张逐笔区只剩 IF($Z=0,"",…) 空壳；第 2～6 行、月份表、隐藏参数
  都套 pause/pause0（make 查）；标题换成 layout.PAUSE_TITLE（条件格式红字），状态格「△ 录入模式中」（登记 ws._pause_exempt）。
隐藏格（layout.SUB_PARAM）：Z1 账户号、Z2 年份、Z3 键（账年键；账户查询选了月份＝账月键）、Z4 笔数（_流 里这个键的条数）、
  Y1 账户名、Y2 子表名（常数；备用/查询＝""）、Y3 月码（子表＝0）、Y4 状态号；U/V 各月扫描窗口（U6/V6＝年初用的「年前」窗口，U19/V19＝全年/选的月）。
私有帮助格（X 列，本模块自己用；s_check 可以读）：X1 最后月码、X2 本表容量、X3 本表明细净额（SUM(G)−SUM(H)）、X4 J4 用的窗口末、
  X5 明细起点余额（＝I6：年初余额；账户查询选了月份＝月初余额）、X6 年份格生效值（查询）、X7 月初余额用的窗口末（查询）；W7..W18 各月月码。
"""
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from layout import (
    SUB_SHEETS, SPARE_SHEETS, SUB_CAP, N_QUERY, SH_QRY, N_FLOW, ACC_TYPES, ACC_STATUS, ATTR_NO,
    SUB_HDR, SUB_OPEN, SUB_R0, SUB_COLS, SUB_HEAD, SUB_WIDTH, SUB_MON, SUB_MON_HDR, SUB_MON_OPEN, SUB_MON_R1, SUB_MON_TOT,
    SUB_HELP, SUB_MON_WIN, SUB_PARAM, SUB_SEL, SUB_KPI, QRY_SEL, QRY_KPI, SUB_FILTER_COLS, RGB_GUESS_FONT,
    RGB_ERR_FILL, RGB_ERR_FONT, PAUSE_TITLE,
    K_YM, K_DAY, FMT_MONEY, lref_z, raw, norm, esc, name2no, yr_eff, yr_bad, ym_eff, win_first, win_last, win_before, win,
    k_aypm, sum_before_year, next_pos, chain, pause, pause0)
from common import (put, title, header, widths, hide, home_link, selector, kpi, cf_flags, cf_neg_red, cf_bg, print_setup,
                    F_TXT, F_TXTB, F_AUTO, F_HELP, F_NOTE, FILL_KPI, FILL_SEL, BD, AC, AL, ALW, AR,
                    ACC, INT, YH)

INT9 = ATTR_NO['内部转账']
MONEY0 = '#,##0.00;[Red]-#,##0.00;'                   # 明细收入/支出：0 不显示（像原日记账，一行只有一栏有数）
FMT_CNT = '0" 笔"'
P = SUB_PARAM                                          # Z1 账户号、Z2 年份、Z3 键、Z4 笔数、Y1 名、Y2 表名、Y3 月码、Y4 状态号
ab = lambda c: f'${c[0]}${c[1:]}'                       # 'Z1' → '$Z$1'
A_, Y_, KEY, CNT = ab(P['账户号']), ab(P['年份']), ab(P['键']), ab(P['笔数'])
NAME, SUBN, YM, ST = ab(P['名']), ab(P['表名']), ab(P['月码']), ab(P['状态号'])
X = {k: P[k] for k in ('最后月码', '容量', '明细净额', '余额窗末', '起点余额', '年份格', '月初窗末')}   # layout.SUB_PARAM 的 X 列
XA = {k: ab(v) for k, v in X.items()}
MC = 'W'                                               # 各月月码（W7..W18）
U, V = SUB_MON_WIN['窗首'], SUB_MON_WIN['窗末']
Z = SUB_HELP['位置']
SC, MO = SUB_COLS, SUB_MON
R_OPEN, R1, RT = SUB_MON_OPEN, SUB_MON_R1, SUB_MON_TOT     # 6 年初；7..18 月；19 全年
HELP_COLS = ('T', U, V, MC, 'X', 'Y', Z)
assert SUB_R0 == 7 and Z == 'Z', '逐笔公式按「第 7 行、Z 列」写的（$Z7），layout 改了这里要跟着改'

# 录入模式：标题（第 1 行）换成 PAUSE_TITLE；状态格（不停，登记 _pause_exempt）
ENTRY_ON = 'P_录入模式=1'
ST_ENTRY = '"△ 录入模式中：本表先不算，录完去基础资料①把「录入模式」清空"'

SUB_TIP = ('从【资金台帐】自动分出来的本账户流水，不用填，也别在这里改（改了跟台帐对不上，右上角会报 ※）：要改就去【资金台帐】改。'
           'H3 填年份看别的年（空＝管理年度）。余额跟台帐一样按上下顺序算，右边月份表按日期算（收入、支出不含内部转账，内部转入/转出单列）。'
           '琥珀色斜体的收费项、收付人＝台帐里空着、按摘要认出来的，对就去台帐选上。点第 5 行的筛选箭头可以按收费项、收付人、摘要筛选。'
           'D3 空着＝本表的账户；把这张表复制给新开的卡用时，在 D3 选那个账户。基础资料①「录入模式」填了「是」时本表先不算（标题变红）。')
# 复核 F4：备用表别改名（首页、核对表的链接按表名跳）；用的是哪个账户看第 1 行标题
SPARE_TIP = ('备用子表：给新开的卡用。在 D3 选账户（下拉，也可以写账户名或别名），这张表就分出那个账户这一年的流水；'
             '表名别改（首页、核对表的链接按表名跳），用的是哪个账户看第 1 行标题。'
             '别的跟账户子表一样：不用填，改动录到【资金台帐】；H3 填年份（空＝管理年度）；余额按台帐上下顺序，右边月份表按日期。')
QRY_TIP = ('任选一个账户查流水：D3 选账户（下拉，也可以写原表名），H3 年份（空＝管理年度），J3 月份（空＝全年；填 9 或 202609）。'
           f'这里最多放 {N_QUERY:,} 笔，一年超过的请选月份。余额按台帐上下顺序，右边月份表按日期（收入、支出不含内部转账）。'
           '不用填，改动录到【资金台帐】。')
# 复核 F10：月份格的出错提示跟输入提示、实际算法一致（1～12 按 H3 那一年）
QRY_MONTH_PROMPT = '空＝全年。填 1～12（H3 那一年的这个月）、202609，或那个月的某一天'
QRY_MONTH_ERROR = '填 1～12（H3 那一年的这个月）、202609 这样的年月，或那个月的某一天'


def _x(f, z='$Z7'):
    """_流 第 z 个位置的某列"""
    return f'INDEX(流_{f},{z})'


def _tx(field, z='$Z7'):
    """_流 位置 z 那一笔在资金台帐的录入原文（去首尾空格）：整列名称 台帐_字段，行号＝段_区首＋取数号−1（layout.lref_z）"""
    return f'TRIM({lref_z(field, z)}&"")'


def _guard(f, z='$Z7'):
    return f'=IF({z}=0,"",{f})'


def _detail_formulas():
    """逐笔各列（第 7 行写法；整列一样，相对引用 $Z7；序号 ROW()−6，不写往下撑的 ROWS($Z$7:Z7)）"""
    ino, pno = _x('项号'), _x('人号')
    to, co = _tx('对方账户'), _tx('归属公司')
    return {
        '序号': _guard(f'ROW()-{SUB_R0 - 1}'),
        '月': _guard(f'MONTH({_x("日期")})'),
        '日': _guard(f'DAY({_x("日期")})'),
        '收费项': _guard(f'IF({ino}>0,INDEX(项_名称,{ino}),{_tx("收费项")})'),
        '收付人': _guard(f'IF({pno}>0,INDEX(人_名称,{pno}),{_tx("收付人")})'),
        '摘要': _guard(_tx('摘要')),
        '收入': _guard(_x('收')),
        '支出': _guard(_x('支')),
        '余额': _guard(_x('余额')),
        '对方归属': _guard(f'IF({to}<>"",{to},IF({co}<>"","归属："&{co},""))'),
        '备注': _guard(_tx('备注')),
    }


def _pos_formulas(qry):
    """Z 列：第 7 行从头 MATCH（链头，pause0：录入模式出 0，整张逐笔区变空壳）；第 8 行起从上一笔往后找（链式，末位 N_FLOW）"""
    if not qry:
        first = pause0(f'=IF({A_}=0,0,IFERROR(MATCH({KEY},流_账年键,0),0))')
        nxt = f'=IF(Z7=0,0,{next_pos("Z7", KEY, "流_账年键", N_FLOW)})'
    else:
        first = pause0(f'=IF({A_}=0,0,IFERROR(IF({YM}=0,MATCH({KEY},流_账年键,0),MATCH({KEY},流_账月键,0)),0))')
        nxt = (f'=IF(Z7=0,0,IF({YM}=0,{next_pos("Z7", KEY, "流_账年键", N_FLOW)},'
               f'{next_pos("Z7", KEY, "流_账月键", N_FLOW)}))')
    return first, nxt


STYLE = {  # 明细各列：字体、对齐、格式
    '序号': (F_AUTO, AC, '0'), '月': (F_TXT, AC, '0'), '日': (F_TXT, AC, '0'), '收费项': (F_TXT, AL, None),
    '收付人': (F_TXT, AL, None), '摘要': (F_TXT, AL, None), '收入': (F_TXT, AR, MONEY0), '支出': (F_TXT, AR, MONEY0),
    '余额': (F_TXTB, AR, FMT_MONEY), '对方归属': (F_TXT, AL, None), '备注': (F_TXT, AL, None),
}


def _style(c, font, align, fmt):
    c.font, c.border, c.alignment = font, BD, align
    if fmt:
        c.number_format = fmt


def _dv_error(ws, cell, text):
    """common.selector 的月份出错提示写死了「管理年度的那个月」：改成本表自己的说法（复核 F10）"""
    for dv in ws.data_validations.dataValidation:
        if cell in str(dv.sqref).split():
            dv.error = text


def build_one(ws, kind, sub=None, orig=None, cap=40):
    """kind：'sub' 账户子表（Y2＝子表名常数）/ 'spare' 备用子表 / 'qry' 账户查询"""
    qry = kind == 'qry'
    end = SUB_R0 + cap - 1
    widths(ws, SUB_WIDTH)
    for c in HELP_COLS:
        ws.column_dimensions[c].width = 9
    # ── 隐藏参数（录入模式：数出 0，名字出 ""）──
    sel_acc, sel_year = SUB_SEL['账户'], SUB_SEL['年份']
    D3, H3 = ab(sel_acc), ab(sel_year)
    J3 = ab(QRY_SEL['月份'])
    ws[P['表名']] = sub if kind == 'sub' else None
    if kind == 'sub':
        by_sub = (f'IF(TRIM({SUBN}&"")="",0,IFERROR(MATCH({esc(norm(SUBN))},账户_子表名规,0),{name2no(SUBN, "账户")}))')
        ws[P['账户号']] = pause0(f'=IF({raw(D3)},{name2no(D3, "账户")},{by_sub})')
    else:
        ws[P['账户号']] = pause0(f'=IF({raw(D3)},{name2no(D3, "账户")},0)')
    if qry:
        ws[X['年份格']] = pause0('=' + yr_eff(H3))
        ws[P['月码']] = pause0('=' + ym_eff(J3, XA['年份格'], '0'))
        ws[P['年份']] = pause0(f'=IF({YM}>0,INT({YM}/100),{XA["年份格"]})')
        ws[P['键']] = pause0(f'=IF({YM}=0,{A_}*10000+{Y_},{A_}*{K_YM}+{YM})')
        ws[P['笔数']] = pause0(f'=IF({A_}=0,0,IF({YM}=0,COUNTIF({win("流_账年键", f"${U}${RT}", f"${V}${RT}")},{KEY}),'
                               f'COUNTIF({win("流_账月键", f"${U}${RT}", f"${V}${RT}")},{KEY})))')
    else:
        ws[P['月码']] = 0
        ws[P['年份']] = pause0('=' + yr_eff(H3))
        ws[P['键']] = pause0(f'={A_}*10000+{Y_}')
        ws[P['笔数']] = pause0(f'=IF({A_}=0,0,COUNTIF({win("流_账年键", f"${U}${RT}", f"${V}${RT}")},{KEY}))')
    ws[P['名']] = pause(f'=IF({A_}=0,"",INDEX(账户_名称,{A_}))')
    ws[P['状态号']] = pause0(f'=IF({A_}=0,0,INDEX(账户_状态号,{A_}))')
    ws[X['最后月码']] = pause0('=YEAR(P_最后日期)*100+MONTH(P_最后日期)')
    ws[X['容量']] = cap
    G, H = SC['收入'], SC['支出']
    ws[X['明细净额']] = pause0(f'=IF({A_}=0,0,ROUND(SUM({G}{SUB_R0}:{G}{end})-SUM({H}{SUB_R0}:{H}{end}),2))')
    ws[X['余额窗末']] = pause0('=' + win_before(f'({Y_}+1)*100+1'))
    opening = f'INDEX(账户_期初,{A_})'
    if qry:
        ws[X['月初窗末']] = pause0(f'=IF({YM}=0,1,{win_before(YM)})')
        wv = lambda nm_: win(nm_, '1', XA['月初窗末'])
        mo_open = (f'ROUND({opening}+SUMIFS({wv("流_净")},{wv("流_账日键")},">="&{A_}*{K_DAY},{wv("流_账日键")},'
                   f'"<"&({A_}*{K_DAY}+DATE(INT({YM}/100),MOD({YM},100),1))),2)')
        ws[X['起点余额']] = pause0(f'=IF({A_}=0,0,IF({YM}=0,${MO["月末余额"]}${R_OPEN},{mo_open}))')
    else:
        ws[X['起点余额']] = pause0(f'=IF({A_}=0,0,${MO["月末余额"]}${R_OPEN})')
    # ── 窗口（U/V）、各月月码（W；ROW()−6，不写往下撑的 ROWS($W$7:W7)）──
    ws[f'{U}{R_OPEN}'] = 1
    ws[f'{V}{R_OPEN}'] = pause0('=' + win_before(f'{Y_}*100+1'))
    for r in range(R1, RT):
        ws[f'{MC}{r}'] = pause0(f'={Y_}*100+ROW()-{R1 - 1}')
        ws[f'{U}{r}'] = pause0('=' + win_first(f'{MC}{r}'))
        ws[f'{V}{r}'] = pause0('=' + win_last(f'{MC}{r}', f'{U}{r}'))
    if qry:
        ws[f'{U}{RT}'] = pause0(f'=IF({YM}=0,{win_first(f"{Y_}*100+1")},{win_first(YM)})')
        ws[f'{V}{RT}'] = pause0(f'=IF({YM}=0,{win_last(f"{Y_}*100+12", f"{U}{RT}")},{win_last(YM, f"{U}{RT}")})')
    else:
        ws[f'{U}{RT}'] = pause0('=' + win_first(f'{Y_}*100+1'))
        ws[f'{V}{RT}'] = pause0('=' + win_last(f'{Y_}*100+12', f'{U}{RT}'))
    for c in HELP_COLS:
        for r in range(1, RT + 1):
            cell = ws[f'{c}{r}']
            if cell.value is not None:
                cell.font = F_HELP
    # ── 第 1 行标题（录入模式＝PAUSE_TITLE，红字）、第 2 行说明 ──
    yr = f'{Y_}&" 年"' + (f'&IF({YM}>0," "&MOD({YM},100)&" 月","")' if qry else '')
    if kind == 'sub':
        nohit = (f'IF({raw(D3)},"【"&TRIM({D3}&"")&"】※ D3 填的账户不在基础资料③：请从下拉选",'
                 f'"【"&{SUBN}&"】※ 基础资料③找不到子表名「"&{SUBN}&"」的账户：子表名列别改")')
        tail = f'IF({raw(D3)},"","（原表：{orig}）")&" · 从【资金台帐】自动分出来，不用填"'
        head = '账户明细'
    elif kind == 'spare':
        nohit = (f'IF({raw(D3)},"※ D3 填的「"&TRIM({D3}&"")&"」不在基础资料③：请从下拉选",'
                 f'"{ws.title}：请在 D3 选账户（给新开的卡用）")')
        tail = '" · 从【资金台帐】自动分出来，不用填"'
        head = '账户明细'
    else:
        nohit = (f'IF({raw(D3)},"账户查询：D3 填的「"&TRIM({D3}&"")&"」不在基础资料③，请从下拉选","账户查询：请在 D3 选账户")')
        tail = '" · 从【资金台帐】查出来，不用填"'
        head = '账户明细查询'
    ws.column_dimensions['A'].width = SUB_WIDTH['A']
    title(ws, f'=IF({ENTRY_ON},"{PAUSE_TITLE}",IF({A_}=0,{nohit},"【"&{NAME}&"】"&{yr}&" {head}"&{tail}))', MO['笔数'],
          {'sub': SUB_TIP, 'spare': SPARE_TIP, 'qry': QRY_TIP}[kind])
    ws.conditional_formatting.add('A1', FormulaRule(formula=[ENTRY_ON], font=Font(name=YH, bold=True, color=RGB_ERR_FONT),
                                                    fill=cf_bg(RGB_ERR_FILL)))
    # ── 第 3 行：回首页｜账户｜信息｜年份｜（月份）｜状态 ──
    home_link(ws, 'A3')
    acc_prompt = ('空着＝本表的账户（' + (sub or '') + '）。把这张表复制给新开的卡用时，在这里选那个账户' if kind == 'sub'
                  else '选账户（下拉；写原表名、旧名也认）')
    selector(ws, 'B3', '账户', sel_acc, dv_formula='账户列表', prompt=acc_prompt)
    ws.merge_cells('B3:C3')
    ws['E3'].fill, ws['E3'].border = FILL_SEL, BD
    ws.merge_cells(f'{sel_acc}:E3')
    cono = f'INDEX(账户_公司号,{A_})'
    types = ','.join(f'"{t}"' for t in ACC_TYPES)
    stats = ','.join(f'"{s}"' for s in ACC_STATUS)
    stop = f'INDEX(账户_停用日,{A_})'
    mdy = lambda d: f'YEAR({d})&"/"&MONTH({d})&"/"&DAY({d})'
    info = (f'=IF({A_}=0,"",IF({cono}>0,INDEX(公司_名称,{cono}),"公司没认出")&" · "&CHOOSE(MAX(1,INDEX(账户_类型号,{A_})),{types})'
            f'&" · "&CHOOSE(MAX(1,MIN(3,{ST})),{stats})&IF(AND({ST}>1,{stop}>0),"（停用 "&{mdy(stop)}&"）","")'
            f'&" · 期初 "&TEXT({opening},"#,##0.00"))')
    put(ws, SUB_KPI['信息'], pause(info), F_NOTE, align=ALW)
    selector(ws, SUB_KPI['年份标签'], '年份', sel_year, kind='year', prompt='空＝管理年度；填 2026、26 或某一天的日期都行')
    if qry:
        selector(ws, QRY_KPI['月份标签'], '月份', QRY_SEL['月份'], kind='month', prompt=QRY_MONTH_PROMPT)
        _dv_error(ws, QRY_SEL['月份'], QRY_MONTH_ERROR)
    st_cell = QRY_KPI['状态'] if qry else SUB_KPI['状态']
    stop_txt = f'"（停用日 "&MONTH({stop})&"/"&DAY({stop})&"）"'
    over = (f'"※ 本年 "&{CNT}&" 笔，本表只放得下 {cap:,} 笔：后面的看【账户查询】选月份"' if not qry else
            f'"※ "&IF({YM}=0,"本年","本月")&" "&{CNT}&" 笔，这里只放得下 {cap:,} 笔："&IF({YM}=0,"请在 J3 选月份","超出的看不到")')
    R_ = MO['月末余额']
    bad_d3 = '"※ D3 填的账户不在基础资料③：请从下拉选"'
    st0 = (f'IF({raw(D3)},{bad_d3},"※ 基础资料③找不到子表名「"&{SUBN}&"」的账户：子表名列别改")' if kind == 'sub'
           else f'IF({raw(D3)},{bad_d3},"请在 D3 选账户")')
    pairs = [
        (ENTRY_ON, ST_ENTRY),                         # 录入模式：只显示这句（别的都不算；这一格登记 _pause_exempt）
        (f'{A_}=0', st0),
        (yr_bad(H3), f'"※ 年份看不懂，按 "&{Y_}&" 年算"'),
    ]
    if qry:
        pairs.append((f'AND({raw(J3)},{YM}=0)', f'"※ 月份看不懂，按 "&{Y_}&" 年全年算"'))
    pairs += [
        (f'{CNT}>{XA["容量"]}', over),
        (f'ROUND({XA["明细净额"]}-(${R_}${RT}-{XA["起点余额"]}),2)<>0',
         '"※ 本表跟台帐对不上：在本表改过？改动要录到【资金台帐】"'),
        (f'{ST}=2', f'"暂停的账户"&IF({stop}>0,{stop_txt},"")'),
        (f'{ST}=3', f'"已销户的账户"&IF({stop}>0,{stop_txt},"")'),
    ]
    put(ws, st_cell, '=' + chain(pairs, '"✓ 跟【资金台帐】一致"'), Font(name=YH, sz=10, bold=True, color='FF375623'), align=ALW)
    ws._pause_exempt = [st_cell]
    ws.merge_cells(f'{st_cell}:{MO["笔数"]}3')
    cf_flags(ws, st_cell, st_cell)
    ws.row_dimensions[3].height = 30
    # ── 第 4 行 KPI ──
    K = QRY_KPI if qry else SUB_KPI
    lab = (lambda y, m: pause(f'=IF({YM}=0,"{y}","{m}")')) if qry else (lambda y, m: y)
    kpi(ws, K['年初标签'], K['年初'], lab('年初余额', '月初余额'), pause(f'=IF({A_}=0,"",${SC["余额"]}${SUB_OPEN})'))
    kpi(ws, K['收入标签'], K['收入'], lab('本年收入\n（不含内部转账）', '本月收入\n（不含内部转账）'),
        pause(f'=IF({A_}=0,"",${MO["收入"]}${RT})'))
    kpi(ws, K['支出标签'], K['支出'], lab('本年支出\n（不含内部转账）', '本月支出\n（不含内部转账）'),
        pause(f'=IF({A_}=0,"",${MO["支出"]}${RT})'))
    kpi(ws, K['余额标签'], K['余额'], pause(f'=IF({Y_}<YEAR(P_最后日期),"年末余额","当前余额")'),
        pause(f'=IF({A_}=0,"",ROUND({opening}+{sum_before_year(win("流_净", "1", XA["余额窗末"]), win("流_账年属月键", "1", XA["余额窗末"]), A_, f"{Y_}+1")},2))'))
    put(ws, K['笔数'], pause(f'={CNT}'), Font(name=YH, sz=11, bold=True, color='FF1F3864'), FILL_KPI, fmt=FMT_CNT, align=AC)
    lb = K['年初标签']
    ws.merge_cells(f'{lb}:{CL(CI(lb[0]) + 1)}{lb[1:]}')
    ws.row_dimensions[4].height = 32
    # ── 第 5 行表头 ──
    header(ws, SUB_HDR, [(SC[k], SUB_HEAD[k]) for k in SC], '子表')
    mon_head = dict(月份='月份', 收入='收入\n（不含内部）', 支出='支出\n（不含内部）', 内部转入='内部转入', 内部转出='内部转出',
                    月末余额='月末余额', 笔数='笔数')
    header(ws, SUB_MON_HDR, [(MO[k], mon_head[k]) for k in MO], '子表')
    # ── 第 6 行「年初余额」──
    r6 = SUB_OPEN
    for k, c in SC.items():
        put(ws, f'{c}{r6}', None, F_TXTB, FILL_KPI, align=AC)
    ws[f'{SC["摘要"]}{r6}'] = pause(f'=IF({YM}=0,"年初余额","月初余额")') if qry else '年初余额'
    ws[f'{SC["余额"]}{r6}'] = pause(f'=IF({A_}=0,"",{XA["起点余额"]})')
    ws[f'{SC["余额"]}{r6}'].number_format = FMT_MONEY
    ws[f'{SC["余额"]}{r6}'].alignment = AR
    # ── 月份表 ──
    Nc, Oc, Pc, Qc, Rc, Sc_, Mc = (MO[k] for k in ('收入', '支出', '内部转入', '内部转出', '月末余额', '笔数', '月份'))
    put(ws, f'{Mc}{R_OPEN}', '年初', F_TXTB, FILL_KPI, align=AC)
    for c in (Nc, Oc, Pc, Qc, Sc_):
        put(ws, f'{c}{R_OPEN}', None, F_TXT, FILL_KPI, fmt=ACC)
    put(ws, f'{Rc}{R_OPEN}', pause(f'=IF({A_}=0,"",ROUND({opening}+'
                                    f'{sum_before_year(win("流_净", f"${U}${R_OPEN}", f"${V}${R_OPEN}"), win("流_账年属月键", f"${U}${R_OPEN}", f"${V}${R_OPEN}"), A_, Y_)},2))'),
        F_TXTB, FILL_KPI, fmt=ACC)
    for r in range(R1, RT):
        m = r - R1 + 1
        w = lambda nm_: win(nm_, f'${U}{r}', f'${V}{r}')
        live = (f'AND({A_}>0,${MC}{r}<={XA["最后月码"]},OR({YM}=0,${MC}{r}={YM}))' if qry
                else f'AND({A_}>0,${MC}{r}<={XA["最后月码"]})')
        kin = k_aypm(A_, Y_, INT9, f'MOD(${MC}{r},100)')
        akey = f'{A_}*{K_YM}+${MC}{r}'
        put(ws, f'{Mc}{r}', f'{m}月', F_TXT, align=AC)
        put(ws, f'{Pc}{r}', pause(f'=IF({live},SUMIF({w("流_账年属月键")},{kin},{w("流_收")}),"")'), F_TXT, fmt=ACC)
        put(ws, f'{Qc}{r}', pause(f'=IF({live},SUMIF({w("流_账年属月键")},{kin},{w("流_支")}),"")'), F_TXT, fmt=ACC)
        put(ws, f'{Nc}{r}', pause(f'=IF({live},ROUND(SUMIF({w("流_账月键")},{akey},{w("流_收")})-{Pc}{r},2),"")'), F_TXT, fmt=ACC)
        put(ws, f'{Oc}{r}', pause(f'=IF({live},ROUND(SUMIF({w("流_账月键")},{akey},{w("流_支")})-{Qc}{r},2),"")'), F_TXT, fmt=ACC)
        prev = (f'IF({YM}=0,{Rc}{r - 1},{XA["起点余额"]})' if qry else f'{Rc}{r - 1}')
        put(ws, f'{Rc}{r}', pause(f'=IF({live},ROUND({prev}+{Nc}{r}-{Oc}{r}+{Pc}{r}-{Qc}{r},2),"")'), F_TXTB, fmt=ACC)
        put(ws, f'{Sc_}{r}', pause(f'=IF({live},COUNTIF({w("流_账月键")},{akey}),"")'), F_TXT, fmt=INT)
    # 第 19 行：全年（账户查询选了月份＝那个月）
    wt = lambda nm_: win(nm_, f'${U}${RT}', f'${V}${RT}')
    yin = lambda col: (f'SUMIFS({wt(col)},{wt("流_账年属月键")},">="&{k_aypm(A_, Y_, INT9, 1)},'
                       f'{wt("流_账年属月键")},"<="&{k_aypm(A_, Y_, INT9, 12)})')
    tot = lambda col: f'SUMIF({wt("流_账年键")},{KEY},{wt(col)})'
    if qry:
        min_ = lambda col: f'SUMIF({wt("流_账年属月键")},{k_aypm(A_, Y_, INT9, f"MOD({YM},100)")},{wt(col)})'
        mtot = lambda col: f'SUMIF({wt("流_账月键")},{KEY},{wt(col)})'
        fin = lambda col: f'IF({YM}=0,{yin(col)},{min_(col)})'
        ftot = lambda col: f'IF({YM}=0,{tot(col)},{mtot(col)})'
        put(ws, f'{Mc}{RT}', pause(f'=IF({YM}=0,"全年","本月")'), F_TXTB, FILL_KPI, align=AC)
    else:
        fin, ftot = yin, tot
        put(ws, f'{Mc}{RT}', '全年', F_TXTB, FILL_KPI, align=AC)
    put(ws, f'{Pc}{RT}', pause(f'=IF({A_}=0,"",ROUND({fin("流_收")},2))'), F_TXTB, FILL_KPI, fmt=ACC)
    put(ws, f'{Qc}{RT}', pause(f'=IF({A_}=0,"",ROUND({fin("流_支")},2))'), F_TXTB, FILL_KPI, fmt=ACC)
    put(ws, f'{Nc}{RT}', pause(f'=IF({A_}=0,"",ROUND({ftot("流_收")}-{Pc}{RT},2))'), F_TXTB, FILL_KPI, fmt=ACC)
    put(ws, f'{Oc}{RT}', pause(f'=IF({A_}=0,"",ROUND({ftot("流_支")}-{Qc}{RT},2))'), F_TXTB, FILL_KPI, fmt=ACC)
    put(ws, f'{Rc}{RT}', pause(f'=IF({A_}=0,"",ROUND({XA["起点余额"]}+{Nc}{RT}-{Oc}{RT}+{Pc}{RT}-{Qc}{RT},2))'), F_TXTB, FILL_KPI, fmt=ACC)
    put(ws, f'{Sc_}{RT}', pause(f'=IF({A_}=0,"",{CNT})'), F_TXTB, FILL_KPI, fmt=INT)
    for c in (Nc, Oc, Pc, Qc, Rc, Sc_):
        ws[f'{c}{RT}'].font = Font(name=YH, sz=10, bold=True, color='FFC00000')
    cf_neg_red(ws, f'{Rc}{R_OPEN}:{Rc}{RT}', f'{Rc}{R_OPEN}')
    # ── 逐笔 ──
    det = _detail_formulas()
    first, nxt = _pos_formulas(qry)
    zfont = F_HELP
    for r in range(SUB_R0, end + 1):
        zc = ws.cell(row=r, column=CI(Z), value=first if r == SUB_R0 else nxt.replace('Z7', f'Z{r - 1}'))
        zc.font = zfont
        for k, col in SC.items():
            c = ws.cell(row=r, column=CI(col), value=det[k].replace('$Z7', f'$Z{r}'))
            _style(c, *STYLE[k])
    # 条件格式：按摘要认出来的收费项（流_项推>0）/收付人（台帐空着、认出来了）琥珀斜体；余额负数红
    dcol, ecol = SC['收费项'], SC['收付人']
    guess = Font(name=YH, italic=True, color=RGB_GUESS_FONT)
    z0 = f'$Z{SUB_R0}'
    ws.conditional_formatting.add(f'{dcol}{SUB_R0}:{dcol}{end}', FormulaRule(
        formula=[f'AND({z0}>0,INDEX(流_项推,{z0})>0)'], font=guess))
    ws.conditional_formatting.add(f'{ecol}{SUB_R0}:{ecol}{end}', FormulaRule(
        formula=[f'AND({z0}>0,{_tx("收付人", z0)}="")'], font=guess))
    cf_neg_red(ws, f'{SC["余额"]}{SUB_R0}:{SC["余额"]}{end}', f'{SC["余额"]}{SUB_R0}')
    hide(ws, *HELP_COLS)
    ws.freeze_panes = f'A{SUB_OPEN}'
    ws.auto_filter.ref = f'{SUB_FILTER_COLS.split(":")[0]}{SUB_HDR}:{SUB_FILTER_COLS.split(":")[1]}{end}'
    print_setup(ws, f'{SUB_HDR}:{SUB_HDR}', landscape=True)


def build(wb, ctx):
    for sub, orig, _cap in SUB_SHEETS:
        build_one(wb[sub], 'sub', sub=sub, orig=orig, cap=SUB_CAP[sub])
    for sp, _cap in SPARE_SHEETS:
        build_one(wb[sp], 'spare', cap=SUB_CAP[sp])
    build_one(wb[SH_QRY], 'qry', cap=N_QUERY)
