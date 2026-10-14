# -*- coding: utf-8 -*-
"""A074 集团多账户资金台帐：把各张表拼成一本：先按顺序建好所有空表（模块之间互相引用的表一定存在），再让各模块往里写；
最后定义名称（v3.1：分段读数区名称 台k_字段 写成隐藏名称）、检查公式、保护查看表。

检查（check_formulas，有一条不过就拦；v3.1 设计 §15.1.9、§15.2.3）：
- Excel 2010 / WPS 规定：不准用的函数、LOOKUP 只准 _流、文字常量 ≤255 字、整条 ≤8000 字、括号 ≤50 层、下拉名称不准进单元格公式、
  用 流_收/流_支 求和要带键列条件、名称要定义过；
- 录入模式（像 v2 的 rpt()）：PAUSE_SHEETS 第 2 行起每个公式最外层套 pause()/pause0()；例外只有 ws._pause_exempt 登记的格、
  子表类（子表、备用、账户查询）第 SUB_R0 行起 A..K 以 =IF($Z{r}=0, 打头、Z 列以 =IF(Z{r−1}=0,0, 打头的格；Z{SUB_R0} 链头必须是 pause0；
- _流（_flow_rules）：原始列只看本行、本段名称 台k_字段、段表本段那一行的链、_表、PAR_SAFE；反查只看段表的链；正常日/余额/顺序引用块表只准写本块那一行；
  问题/提醒是 ≤FLOW_MERGE_FMAX 字、只引用本行的短合并列；配对/期人键 pause 打头；段号写常数；每格 ≤FLOW_FMAX 字；
  不准出现「资金台帐!」「$1048576」（段表除外）、不准引用 PAR_SEGCHK；
- 段表（_seg_rules）：链只引用锚列和段表自己的链；叶子（错位、有字）可以再引用本段的 台k_字段；
  全书只有 _参 的 P_错位、P_笔数差 准引用 段_错位、段_有字；
- WPS 存盘体检（wps_check）：WPS_STRICT 里的表，同一列连续 ≥3 格同写法的公式出现「表名!」或往下撑的区域就报错
  （白名单：_流 段表各列、_流 第 1 段）；_表、_流、子表类里一律不准直接写「资金台帐!」「基础资料!」（用名称）；其余表只打印会被拆的格数、字节。
再查环（cyclecheck）。"""
import importlib
import re
from collections import defaultdict
import openpyxl
from openpyxl.styles import Protection
from openpyxl.utils import range_boundaries
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.protection import SheetProtection
from layout import *
from openpyxl.utils import column_index_from_string as CI, get_column_letter as CL
from common import print_setup, F_HELP
from share import normalize, wps_split, xml_bytes
import data

_V31 = 'FLOW_SEG' in globals()                     # v3.1 契约（git 4575e4b 的 v2 layout 没有这些：只跑 v2 的检查，给「没把 v2 弄坏」做对照）
_G = globals()

# 模块顺序：录入/基础资料 → 隐藏计算 → 子表 → 报表 → 首页/年末余额 → 核对/记录/说明（每个 s_*.py 提供 build(wb, ctx)）
MODULES = ['s_input', 's_mirror', 's_sub', 's_rep1', 's_rep2', 's_home', 's_check']
KEEP_OPEN = {RGB_IN, RGB_SEL}                     # 淡黄手填格、黄色选择格：保护时不锁
GUARDS = tuple(_G.get('PAUSE_GUARDS', (PAUSE_GUARD,)))   # v3.1：pause / pause0 两种打头
SUBLIKE = (set(ALL_SUBS) | {SH_QRY}) if _V31 else set()  # 子表类：链头 Z{SUB_R0}，逐笔区 IF($Z=0,"",…) 空壳
WPS_WARN_MB = 55                                   # build：「WPS 存后估计」超了报警（设计 §15.7-6）

BANNED = re.compile(r'(?<![A-Z0-9_.])(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|LET|LAMBDA|MAXIFS|MINIFS|IFS|SWITCH|'
                    r'TEXTJOIN|CONCAT|INDIRECT|OFFSET|TODAY|NOW|RAND|RANDBETWEEN|AGGREGATE|TEXTSPLIT|PIVOTBY|TRIMRANGE|'
                    r'SHEETS|SHEET|ISFORMULA|FORMULATEXT|IFNA|DAYS)\(')
LOOKUP_RE = re.compile(r'(?<![A-Z0-9_.])LOOKUP\(')        # 只准 _流 的规则列用 LOOKUP(1,0/(…),…)
LOOKUP_OK = {H_FLOW}
# 公式里的引用：[表!]格子/区域、整列、名称（字符串外）。a1/a2＝行号前有没有 $
_RX = re.compile(
    r"(?<![\w.$'#])"
    r"(?:(?:'(?P<qs>(?:[^']|'')+)'|(?P<us>[^\W\d][\w.]*))!)?"
    r"(?:(?P<c1>\$?[A-Z]{1,3})(?P<a1>\$?)(?P<r1>\d+)(?::(?P<c2>\$?[A-Z]{1,3})(?P<a2>\$?)(?P<r2>\d+))?"
    r"|(?P<k1>\$?[A-Z]{1,3}):(?P<k2>\$?[A-Z]{1,3})"
    r"|(?P<nm>[^\W\d][\w.]*))"
    r"(?![\w.(!])")
_NOT_NAMES = {'TRUE', 'FALSE'}
_ERRLIT = re.compile(r'#(?:N/A|REF!|VALUE!|DIV/0!|NUM!|NAME\?|NULL!)')


def _outside(f):
    """字符串外面的部分（"…" 里的字换成空）"""
    return '""'.join(f.split('"')[0::2]) if '"' in f else f


def _refs(f):
    """公式 → (区域列表 [(表或 None, 列1, 列2, 行1, 行2, 行1绝对, 行2绝对)], 名称列表)。整列的行记 (1, 1048576)"""
    refs, names = [], []
    for m in _RX.finditer(_ERRLIT.sub(' ', _outside(f))):
        g = m.group
        sh = g('qs').replace("''", "'") if g('qs') is not None else g('us')
        if g('nm') is not None:
            if sh is None and g('nm').upper() not in _NOT_NAMES:
                names.append(g('nm'))
            continue
        if g('c1') is not None:
            c1, r1, a1 = CI(g('c1').lstrip('$')), int(g('r1')), bool(g('a1'))
            if g('c2') is not None:
                c2, r2, a2 = CI(g('c2').lstrip('$')), int(g('r2')), bool(g('a2'))
            else:
                c2, r2, a2 = c1, r1, a1
            refs.append((sh, min(c1, c2), max(c1, c2), min(r1, r2), max(r1, r2), a1, a2))
        else:
            c1, c2 = sorted((CI(g('k1').lstrip('$')), CI(g('k2').lstrip('$'))))
            refs.append((sh, c1, c2, 1, 1048576, True, True))
    return refs, names


def _balanced(f):
    depth, q_ = 0, False
    for ch in f:
        if ch == '"':
            q_ = not q_
        elif not q_:
            depth += (ch == '(') - (ch == ')')
            if depth < 0:
                return False
    return depth == 0 and not q_


def _depth(f):
    """括号最大深度（近似函数嵌套层数）"""
    d = m = 0
    q_ = False
    for ch in f:
        if ch == '"':
            q_ = not q_
        elif not q_:
            if ch == '(':
                d += 1
                m = max(m, d)
            elif ch == ')':
                d -= 1
    return m


_SUMIF_RE = re.compile(r'(?<![A-Z0-9_.])(SUMIFS?)\(')
_KEYNAME = re.compile(r'流_(\w*键|月码|有效)')


def _args(f, i):
    """f[i] 是左括号：返回 (顶层参数列表, 右括号位置)"""
    depth, q_, cur, out = 0, False, '', []
    for j in range(i, len(f)):
        ch = f[j]
        if ch == '"':
            q_ = not q_
        if not q_:
            if ch == '(':
                depth += 1
                if depth == 1:
                    continue
            elif ch == ')':
                depth -= 1
                if depth == 0:
                    out.append(cur)
                    return out, j
            elif ch == ',' and depth == 1:
                out.append(cur)
                cur = ''
                continue
        cur += ch
    return out, len(f)


def _sumif_needs_key(f):
    """用 流_收/流_支 求和的 SUMIF(S) 条件里必须有键列、流_月码 或 流_有效（无效行这几列是 0）→ 返回违规的片段"""
    bad = []
    for m in _SUMIF_RE.finditer(f):
        args, _e = _args(f, m.end() - 1)
        if m.group(1) == 'SUMIF':
            if len(args) < 3:
                continue
            summ, crit_ranges = args[2], [args[0]]
        else:
            summ, crit_ranges = args[0], args[1::2]
        if re.search(r'流_(收|支)(?![\w])', summ) and not any(_KEYNAME.search(r) for r in crit_ranges):
            bad.append(f[m.start():m.start() + 80])
    return bad


def _formulas(ws):
    """表里的公式格（按列、行排好）：[(行, 列号, 公式, 写法)]；写法＝share.normalize（相对行号换成「差几行」，同一列一样＝同一个公式）。
       一张表只算一次（ws._a074_f），各项检查共用"""
    got = getattr(ws, '_a074_f', None)
    if got is None:
        got = [(r, c, cell._value) for (r, c), cell in ws._cells.items()
               if cell.data_type == 'f' and isinstance(cell._value, str) and cell._value.startswith('=')]
        got.sort(key=lambda t: (t[1], t[0]))
        got = [(r, c, f, normalize(f, r)) for r, c, f in got]
        ws._a074_f = got
    return got


def _text_issues(ws_title, f, known):
    """只看公式文字的检查（同一列同写法的格只查一次）"""
    hit = []
    outside = _outside(f).upper()
    hit += BANNED.findall(outside) + (['_xlfn'] if '_XLFN.' in outside else [])
    if ws_title not in LOOKUP_OK and LOOKUP_RE.search(outside):
        hit.append('LOOKUP 只准在 _流 的摘要规则列用')
    long_lit = [len(t) for t in f.split('"')[1::2] if len(t) > 255]       # 公式里的文字常量 Excel 最多 255 字
    if long_lit:
        hit.append(f'文字常量超过 255 字 {long_lit}')
    _r, names = _refs(f)
    dv_use = sorted({n for n in names if n in DV_NAMES})
    if dv_use:
        hit.append(f'引用了下拉名称 {dv_use}（OFFSET 易失）')
    unk = sorted({n for n in names if n.upper() not in known})
    if unk:
        hit.append(f'名称没定义 {unk[:5]}')
    depth = _depth(f)
    if depth > 50:
        hit.append(f'嵌套 {depth} 层')
    if '流_收' in f or '流_支' in f:
        nokey = _sumif_needs_key(f)
        if nokey:
            hit.append(f'流_收/流_支 求和没有键列条件：{nokey[0]}')
    if not _balanced(f):
        hit.append('括号或引号不配对')
    if len(f) > 8000:
        hit.append(f'整条 {len(f)} 字（Excel 最多 8000）')
    return hit


def _exempt(ws):
    """ws._pause_exempt：格子或区域（'H1'、'A3:F3'）→ {(行, 列号)}"""
    out = set()
    for it in getattr(ws, '_pause_exempt', ()) or ():
        it = str(it).replace('$', '')
        c1, r1, c2, r2 = range_boundaries(it if ':' in it else f'{it}:{it}')
        for r in range(r1, r2 + 1):
            for c in range(c1, c2 + 1):
                out.add((r, c))
    return out


def _pause_issue(ws_title, r, c, f, exempt):
    """录入模式（v3.1）/报表暂停（v2）：PAUSE_SHEETS 第 2 行起的公式要 pause()/pause0() 打头"""
    if ws_title not in PAUSE_SHEETS or r < 2:
        return None
    sub = ws_title in SUBLIKE and r >= SUB_R0
    zc = CI(SUB_HELP['位置']) if _V31 else 0
    if sub and c == zc and r == SUB_R0:
        return None if f.startswith(GUARDS[-1]) else f'子表链头 {SUB_HELP["位置"]}{SUB_R0} 要套 layout.pause0()（录入模式出 0）'
    if f.startswith(GUARDS) or (r, c) in exempt:
        return None
    if sub:
        if c <= CI(SUB_FILTER_COLS.split(':')[1]) and re.match(rf'=IF\(\$?{SUB_HELP["位置"]}{r}=0,', f):
            return None
        if c == zc and re.match(rf'=IF\(\$?{SUB_HELP["位置"]}{r - 1}=0,0,', f):
            return None
    return ('录入模式要停的表：公式最外层要套 layout.pause()/pause0()（不停的格登记 ws._pause_exempt）' if _V31
            else '报表公式要套 layout.rpt()（报表暂停）')


# ═════════════════════════════ _流 / 段表 规则（v3.1） ═════════════════════════════
def _name_kind():
    """名称 → 种类：par / seg(k) / seg_chain / seg_leaf / blk / flow / tab（_表、基础资料、公式名称）/ other（台帐整列、报表、核对…）"""
    kinds = {}
    seg_k = {}
    for k in range(1, N_SEG + 1):
        for fld in LEDGER_COLS:
            seg_k[seg_name(fld, k)] = k
    for nm, ref in NAMES.items():
        if nm in seg_k:
            kinds[nm] = ('seg', seg_k[nm])
        elif nm.startswith('P_'):
            kinds[nm] = ('par', nm[2:])
        elif nm.startswith('段_'):
            kinds[nm] = ('seg_leaf' if nm[2:] in SEG_LEAF else 'seg_chain', nm[2:])
        elif nm.startswith('块_'):
            kinds[nm] = ('blk', None)
        elif nm.startswith('流_'):
            kinds[nm] = ('flow', nm[2:])
        elif ref.startswith(f'{H_TAB}!') or ref.startswith(f'{q(SH_BASE)}!') or '!' not in ref:
            kinds[nm] = ('tab', None)                  # _表、基础资料各块、规项R_ 这种公式名称
        else:
            kinds[nm] = ('other', None)
    return kinds


def _flow_rules(ws):
    """v3.1 _流 规则（设计 §15.1.9）：按「列 × 段 × 同写法」只查一次；每条毛病报一次（第一格＋同样的格数）"""
    if not _V31:
        return _flow_rules_v2(ws)
    kinds = _name_kind()
    fld = {CI(c): f for c, f, _ in FLOW_COLS}
    col_of = {f: CI(c) for c, f, _ in FLOW_COLS}
    seg_ci = {CI(c): k for k, c in FLOW_SEG.items()}
    chain_ci = {CI(FLOW_SEG[k]) for k in SEG_CHAIN}
    leaf_ci = {CI(FLOW_SEG[k]) for k in SEG_LEAF}
    blk_lo, blk_hi = CI(FLOW_BLK['净']), CI(FLOW_BLK_LAST)
    flow_hi = CI(FLOW_LAST)
    merge_ok = {col_of[f] for f in ('取数号', '有效', '问本', '顺序', '提本', '提醒') if f in col_of}
    segchk = {f'P_{k}' for k in PAR_SEGCHK}
    bad = defaultdict(list)                    # 毛病 → [格子]
    cache = {}

    def per_row(r, c, f):
        k = slot_of_row(r)[0]
        name = fld[c]
        out = []
        refs, names = _refs(f)
        own_or_chain = []
        for (sh, c1, c2, r1, r2, _a1, _a2) in refs:
            if sh is not None and sh != H_FLOW:
                continue                                         # 跨表：另由「不准直接写别的表」查
            if any(c1 <= x <= c2 for x in leaf_ci):
                out.append('引用了段表叶子（段_错位/有字）')
            if c1 <= blk_hi and c2 >= blk_lo:
                if name in FLOW_CHAIN:
                    if not (r1 == r2 == k + 1):
                        out.append(f'块表只准写本块那一行（第 {k + 1} 行）')
                else:
                    out.append('引用了块表')
            own_or_chain.append((c1, c2, r1, r2))
        for nm in names:
            kd = kinds.get(nm, ('?', None))
            if kd[0] == 'seg_leaf':
                out.append(f'引用了段表叶子 {nm}')
            if kd[0] == 'blk':
                out.append(f'引用了块表名称 {nm}' + ('（正常日要直接写本块「前最晚」那一格，会成环）' if name == '正常日' else ''))
            if kd[0] == 'seg' and kd[1] != k:
                out.append(f'用了别的段的名称 {nm}（本行是第 {k} 段）')
        if name in FLOW_RAW:
            for nm in names:
                kd = kinds.get(nm, ('?', None))
                if kd[0] == 'par' and kd[1] not in PAR_SAFE:
                    out.append(f'原始列引用了 {nm}（只准 PAR_SAFE）')
                elif kd[0] in ('flow', 'other'):
                    out.append(f'原始列引用了 {nm}（只看本行、本段名称、_表）')
            for (c1, c2, r1, r2) in own_or_chain:
                if r1 == r2 == r and c2 <= flow_hi:
                    continue
                if r1 == r2 == k + 1 and all(x in chain_ci for x in range(c1, c2 + 1)):
                    continue
                out.append('原始列只准引用本行、段表本段那一行的链')
        elif name in FLOW_STATIC:
            for nm in names:
                kd = kinds.get(nm, ('?', None))
                if not (kd[0] == 'seg_chain' and kd[1] != '夹缝'):
                    out.append(f'反查只准引用段表的链，不准 {nm}')
            for (c1, c2, r1, r2) in own_or_chain:
                if not all(x in chain_ci for x in range(c1, c2 + 1)) or c1 == CI(FLOW_SEG['夹缝']):
                    out.append('反查只准引用段表的链')
        elif name in FLOW_MERGE:
            if len(f) > FLOW_MERGE_FMAX:
                out.append(f'合并列 {len(f)} 字（≤{FLOW_MERGE_FMAX}）')
            if names:
                out.append(f'合并列不准用名称 {sorted(set(names))[:3]}')
            for (c1, c2, r1, r2) in own_or_chain:
                if not (r1 == r2 == r and c1 == c2 and c1 in merge_ok):
                    out.append('合并列只准引用本行的 取数号/有效/问本/顺序/提本/提醒')
        if name in FLOW_GATED and not f.startswith(GUARDS):
            out.append('配对/期人键 要 pause()/pause0() 打头（录入模式停）')
        if name in FLOW_CONST:
            out.append('段号要写常数，不是公式')
        return out

    def seg_cell(r, c, f):
        k = r - 1
        out = []
        refs, names = _refs(f)
        is_chain = c in chain_ci
        for nm in names:
            kd = kinds.get(nm, ('?', None))
            if is_chain or not (kd[0] == 'seg' and kd[1] == k):
                out.append(f'段表{"链" if is_chain else "叶子"}不准引用名称 {nm}' + ('' if is_chain else f'（只准本段 台{k}_字段）'))
        for (sh, c1, c2, r1, r2, _a1, _a2) in refs:
            if sh is None or sh == H_FLOW:
                ok = chain_ci if is_chain else set(seg_ci)
                if not all(x in ok for x in range(c1, c2 + 1)):
                    out.append('段表' + ('链只准引用锚列和段表的链' if is_chain else '叶子只准引用锚列、段表、本段名称'))
            elif sh == SH_LEDGER:
                if not (c1 == c2 == CI(LEDGER_ANCHOR)):
                    out.append(f'段表只准直接引用锚列 {LEDGER_ANCHOR}')
            else:
                out.append(f'段表引用了 {sh}!')
        return out

    allf = _formulas(ws)
    if not any(c in fld and r >= 2 for r, c, _f, _k in allf):
        return []                                               # 还没有 _流（只生成了别的模块）
    for r, c, f, nk in allf:
        where = f'{CL(c)}{r}'
        lim = FLOW_FMAX_EXC.get(fld.get(c), FLOW_FMAX)
        if len(f) > lim:
            bad[f'公式 {len(f)} 字，超过 FLOW_FMAX {lim}'].append(where)
        outside = _outside(f)
        in_seg = c in seg_ci and 2 <= r <= N_SEG + 1
        if not in_seg and (f'{q(SH_LEDGER)}!' in outside or '$1048576' in outside):
            bad['_流 里不准出现「资金台帐!」「$1048576」（段表除外；取台帐用 layout.sref）'].append(where)
        if any(p in outside for p in segchk) and re.search('|'.join(rf'(?<![\w.]){re.escape(p)}(?![\w])' for p in segchk), outside):
            bad['_流 不准引用 PAR_SEGCHK（P_错位/笔数差/没算/锚坏/可录：只给首页横幅、台帐汇总条、核对表）'].append(where)
        if in_seg:
            key = ('seg', c, r)
            issues = seg_cell(r, c, f)
        elif c in fld and 2 <= r <= N_FLOW + 1:
            key = (c, slot_of_row(r)[0], nk)
            issues = cache.get(key)
            if issues is None:
                issues = cache[key] = per_row(r, c, f)
        else:
            issues = []
        for t in dict.fromkeys(issues):
            bad[f'{fld.get(c) or FLOW_SEG_NAME.get(c, CL(c))}：{t}'].append(where)
    for f_ in FLOW_CONST:                                       # 段号：常数＝本行第几段
        c = CI(FLOW[f_])
        for r in range(2, N_FLOW + 1 + 1):
            v = ws.cell(row=r, column=c).value if (r, c) in ws._cells else None
            if isinstance(v, str) and v.startswith('='):
                continue                                         # 上面已报
            if v != slot_of_row(r)[0]:
                bad[f'{f_} 要写常数＝本行第几段'].append(f'{FLOW[f_]}{r}')
    return [f'{ws.title}!{v[0]}（同样 {len(v)} 格）：{k}' for k, v in bad.items()]


FLOW_SEG_NAME = {CI(c): f'段_{k}' for k, c in _G.get('FLOW_SEG', {}).items()}


def _flow_rules_v2(ws):
    """v2：_流 原始列（A..月码）只准引用 PAR_SAFE 里的 P_，不准引用块表；正常日 不准写 INDEX(块_…)"""
    bad = []
    raw_last = CI(FLOW_RAW_LAST)
    for r, c, v, _k in _formulas(ws):
        if r < 2 or c > CI(FLOW_LAST):
            continue
        outside = ''.join(v.split('"')[0::2])
        if c <= raw_last:
            ps = {p for p in re.findall(r'(?<![\w.])P_(\w+)', outside)} - set(PAR_SAFE)
            if ps or '块_' in outside:
                bad.append(f'{ws.title}!{CL(c)}{r}: 原始列引用了 {sorted(ps) or "块表"}')
        if CL(c) == FLOW['正常日'] and '块_' in outside:
            bad.append(f'{ws.title}!{CL(c)}{r}: 正常日 要直接写本块「前最晚」那一格，不能 INDEX(块_…)（会成环）')
    return bad


def _leaf_rules(wb):
    """全书只有 _参 的 P_错位、P_笔数差 准引用 段_错位、段_有字（叶子；别处引用就整本又变脏）"""
    if not _V31:
        return []
    ok = {(H_PAR, PAR['错位']), (H_PAR, PAR['笔数差'])}
    leaf_names = {f'段_{k}' for k in SEG_LEAF}
    leaf_ci = {CI(FLOW_SEG[k]) for k in SEG_LEAF}
    seg_ci = {CI(c) for c in FLOW_SEG.values()}
    bad = []
    for ws in wb.worksheets:
        for r, c, f, _k in _formulas(ws):
            if '段_' not in f and '!' not in f:
                continue
            refs, names = _refs(f)
            hit = sorted(set(names) & leaf_names)
            for (sh, c1, c2, _r1, _r2, _a, _b) in refs:
                if (sh == H_FLOW or (sh is None and ws.title == H_FLOW and c not in seg_ci)) and \
                        any(c1 <= x <= c2 for x in leaf_ci):
                    hit.append('段表叶子那几列')
            if hit and (ws.title, f'{CL(c)}{r}') not in ok:
                bad.append(f'{ws.title}!{CL(c)}{r}: 引用了 {hit}（只准 _参 的 P_错位、P_笔数差）')
    return bad[:40]


# ═════════════════════════════ WPS 存盘体检（v3.1） ═════════════════════════════
def _wps_white(title, c, r):
    """白名单：_流 段表各列、_流 第 1 段（第 2～SEG_SLOTS+1 行）"""
    if title != H_FLOW or not _V31:
        return False
    return c in FLOW_SEG_NAME or 2 <= r <= SEG_SLOTS + 1


def wps_check(wb):
    """→ (报错列表, {表名: (会被拆的格数, 字节)})。同一列连续 ≥3 格同写法（R1C1 一样）的公式，
       出现「表名!」或往下撑的区域（一端绝对行、一端相对行）WPS 存盘就一格一条全文写（设计 §15.1.9）"""
    strict = set(_G.get('WPS_STRICT', ()))
    no_direct = ({H_FLOW, H_TAB} | SUBLIKE) if _V31 else set()
    direct_pat = re.compile(rf"(?<![\w.']){re.escape(q(SH_LEDGER))}!|(?<![\w.']){re.escape(q(SH_BASE))}!")
    errs, info = [], {}
    for ws in wb.worksheets:
        title = ws.title
        bycol = defaultdict(list)
        for r, c, f, k in _formulas(ws):
            bycol[c].append((r, f, k))
        n = b = 0
        for c, lst in bycol.items():
            i = 0
            keys = [k for _r, _f, k in lst]
            lst = [(r, f) for r, f, _k in lst]
            while i < len(lst):
                j = i + 1
                while j < len(lst) and lst[j][0] == lst[j - 1][0] + 1 and keys[j] == keys[i]:
                    j += 1
                if j - i >= 3:
                    xs, mx = wps_split(lst[i][1])
                    if xs or mx:
                        n += j - i
                        b += sum(xml_bytes(f) for _r, f in lst[i:j])
                        out = [r for r, _f in lst[i:j] if not _wps_white(title, c, r)]
                        if title in strict and out:
                            why = '直接写了「表名!」' if xs else '往下撑的区域（$X$3:X3 这种）'
                            errs.append(f'{title}!{CL(c)}{lst[i][0]}:{CL(c)}{lst[j - 1][0]}（{len(out)} 格）{why}：{lst[i][1][:120]}')
                i = j
            if title in no_direct and not (title == H_FLOW and c in FLOW_SEG_NAME):
                hit = [r for r, f in lst if direct_pat.search(_outside(f))]
                if hit:
                    errs.append(f'{title}!{CL(c)}{hit[0]}（这一列 {len(hit)} 格）: 公式里不准直接写「{SH_LEDGER}!」「{SH_BASE}!」'
                                f'（用名称 台帐_字段 / 基_块_字段 / 台k_字段）')
        info[title] = (n, b)
    return errs, info


def entry_stats(wb):
    """录入模式下各表还会算的公式（pause/pause0 打头的、子表逐笔空壳、_流 配对/期人键 不算）→ {表: (公式数, 照算的, 空壳)}"""
    out = {}
    gated = {CI(FLOW[f]) for f in _G.get('FLOW_GATED', ())}
    zc = CI(SUB_HELP['位置']) if _V31 else 0
    for ws in wb.worksheets:
        tot = live = shell = 0
        for r, c, f, _k in _formulas(ws):
            tot += 1
            if f.startswith(GUARDS) or (ws.title == H_FLOW and c in gated and 2 <= r):
                continue
            if ws.title in SUBLIKE and r >= SUB_R0 and (re.match(rf'=IF\(\$?{SUB_HELP["位置"]}\d+=0,', f) or c == zc):
                shell += 1
                continue
            live += 1
        if tot:
            out[ws.title] = (tot, live, shell)
    return out


def check_formulas(wb):
    known = {n.upper() for n in set(NAMES) | set(DV_NAMES) | set(wb.defined_names.keys())}   # Excel 的名称不分大小写
    agg = {}                                         # (表, 列, 毛病) → [第一格, 格数, 公式]：同一列同样的毛病只报一条
    cache = {}
    for ws in wb.worksheets:
        ex = _exempt(ws)
        for r, c, v, nk in _formulas(ws):
            key = (ws.title, c, nk)
            hit = cache.get(key)
            if hit is None:
                hit = cache[key] = _text_issues(ws.title, v, known)
            p = _pause_issue(ws.title, r, c, v, ex)
            if hit or p:
                k = (ws.title, c, tuple(hit + ([p] if p else [])))
                if k in agg:
                    agg[k][1] += 1
                else:
                    agg[k] = [f'{CL(c)}{r}', 1, v]
    bad = [f'{t}!{a}（这一列同样 {n} 格）: {list(h)} len={len(v)} {v[:160]}' for (t, _c, h), (a, n, v) in agg.items()]
    if H_FLOW in wb.sheetnames:
        bad += _flow_rules(wb[H_FLOW])
    bad += _leaf_rules(wb)
    werr, winfo = wps_check(wb)
    bad += werr
    assert not bad, f'公式检查没过 {len(bad)} 条：\n' + '\n'.join(bad[:40])
    for ws in wb.worksheets:                         # 有效性的提示、出错文字 Excel 最多 255 字
        for dv in ws.data_validations.dataValidation:
            for t in (dv.prompt, dv.error, dv.promptTitle, dv.errorTitle):
                assert t is None or len(t) <= 255, f'{ws.title} 有效性文字超过 255 字：{t[:60]}…（{len(t)}）'
            f1 = dv.formula1 or ''
            assert not f1.startswith('='), f'{ws.title} 有效性 formula1 不要写开头的 =：{f1[:60]}'
            if dv.type == 'list' and f1.startswith('"'):
                assert len(f1) <= 257, f'{ws.title} 下拉直接写的清单超过 255 字'
    return winfo


def _no_emoji(wb):
    """💡 在老电脑的 Excel / WPS 上可能显示成方框：统一换成 ★"""
    for ws in wb.worksheets:
        for cell in ws._cells.values():
            v = cell._value
            if isinstance(v, str) and '💡' in v:
                cell.value = v.replace('💡 ', '★ ').replace('💡', '★ ')
        for dv in ws.data_validations.dataValidation:
            for a in ('prompt', 'error', 'promptTitle', 'errorTitle'):
                t = getattr(dv, a)
                if t and '💡' in t:
                    setattr(dv, a, t.replace('💡 ', '★ ').replace('💡', '★ '))


def define_names(wb):
    """NAMES、下拉名称；v3.1 的分段读数区名称（NAMES_HIDDEN，台k_字段 约 1,100 个）写成隐藏名称（名称管理器里不显示）"""
    hidden = set(_G.get('NAMES_HIDDEN', ()))
    for nm, ref in NAMES.items():
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref, hidden=True if nm in hidden else None)
    for nm, ref in DV_NAMES.items():
        wb.defined_names[nm] = DefinedName(nm, attr_text=ref)


def _fit_widths(wb):
    for ws in wb.worksheets:
        if ws.sheet_state != 'visible':
            continue
        need = {}
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 400)):
            for c in row:
                if c.value is None:
                    continue
                fmt = c.number_format or ''
                sz = (c.font.sz or 10) if c.font is not None else 10
                w = 12.5 if ('yy' in fmt or 'm"' in fmt) else (13.5 if '#,##0.00' in fmt else (11 if '#,##0' in fmt else 0))
                if w:
                    w = w * max(1.0, sz / 10.0) + (1.5 if (c.font is not None and c.font.b) else 0)
                    need[c.column_letter] = max(need.get(c.column_letter, 0), w)
        for col, w in need.items():
            cd = ws.column_dimensions[col]
            if not cd.hidden and (cd.width or 8.43) < w:
                cd.width = w


def build(only=None, scenario=None, ctx=None):
    errs = self_check()
    assert not errs, 'layout 自检没过：\n' + '\n'.join(errs)
    ctx = ctx if ctx is not None else data.build_ctx()
    if scenario:
        scenario(ctx)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for _g, color, names in GROUPS:
        for n in names:
            ws = wb.create_sheet(n)
            ws.sheet_properties.tabColor = color[2:]
    for n in HIDDEN:
        wb.create_sheet(n)
    define_names(wb)
    for m in MODULES:
        if only and m not in only:
            continue
        try:
            mod = importlib.import_module(m)
        except ModuleNotFoundError as e:
            if e.name == m:
                print(f'   （模块 {m} 还没写，跳过）')
                continue
            raise
        mod.build(wb, ctx)
    for n in HIDDEN:
        wb[n].sheet_state = 'hidden'
        for row in wb[n].iter_rows(min_row=1, max_row=1):
            for c in row:
                c.font = F_HELP
    for grp, color, names in GROUPS:
        for n in names:
            ws = wb[n]
            ws.sheet_properties.tabColor = color[2:]
            if grp in VIEW_GROUPS and not getattr(ws, '_no_protect', False):
                for row in ws.iter_rows():
                    for c in row:
                        if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb in KEEP_OPEN:
                            c.protection = Protection(locked=False)
                ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                                formatCells=False, selectLockedCells=False, selectUnlockedCells=False)
            if not ws.print_title_rows and grp not in ('首页',) and not getattr(ws, '_keep_print', False):
                print_setup(ws, None, landscape=True)
    _no_emoji(wb)
    bad = [nm for nm in wb.defined_names if not re.fullmatch(r'[A-Za-z_\\一-鿿][\w.一-鿿]*', nm)]
    assert not bad, f'名称里有 Excel 不认的字符：{bad[:10]}'
    winfo = check_formulas(wb)
    soft = {k: v for k, v in winfo.items() if v[0] and k not in set(_G.get('WPS_STRICT', ()))}
    if _V31 and soft:
        print('   WPS 存盘会拆开的（不拦）：' + '、'.join(f'{k} {n:,} 格 {b / 1e6:.1f} MB' for k, (n, b) in
                                              sorted(soft.items(), key=lambda kv: -kv[1][1])[:10]))
    import cyclecheck                                  # 静态环检测（Excel 把 INDEX/SUMIF 的整片区域都当前驱）
    cyc = cyclecheck.find_cycles(wb)
    if cyc:
        print(f'✗ 公式有环 {len(cyc)} 处（Excel 打开会报循环引用），前 10 个：\n' + cyclecheck.fmt(cyc[:10]))
    assert not cyc, f'公式有环 {len(cyc)} 处（见上面打印的前 10 个）'
    _fit_widths(wb)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 0
    wb.worksheets[0].sheet_view.tabSelected = True
    wb.calculation.fullCalcOnLoad = False
    wb.calculation.calcId = 191029
    wb._a074_wps = winfo                               # build 打印用
    wb._a074_entry = entry_stats(wb)
    wb._a074_ctx = ctx
    for ws in wb.worksheets:
        ws.__dict__.pop('_a074_f', None)
    return wb
