# -*- coding: utf-8 -*-
"""报表一（s_rep1）：【资金日报】（参考模板汇报表的三件套）、【月度汇报表】、【公司汇总】。设计见 设计.md §5.2～§5.4。

口径（设计 §1.3）：
  账户口径（资金流）：收入＝收入栏（不含内部转账）、支出＝支出栏（不含内部转账）、内部转入、内部转出分列；
    期末＝期初＋收入－支出＋内部转入－内部转出。资金日报 ① 照参考模板是「收款/付款（含内部转入/转出）」。
  经营口径（收支分析）：经营收入＝属性「收入」行的 Σ(收－支)；经营支出＝属性「支出」行的 Σ(支－收)；投资建设＝Σ(支－收)；
    往来流入/流出＝往来行的收入栏/支出栏；待分类、内部净额＝Σ(收－支)。
  余额一律按「日期」算（期初＋日期在那天以前的收－支），跟台帐上下顺序无关。

写法：
  - 只引用定义名称（_流、_表、_参、块_）和本表格子；名单按序号 INDEX(账户_名称,k)、INDEX(公司_名称,c) 取；两级清单用 账两级_候选键。
  - 扫 _流 一律配窗口（layout.win_*，窗口首/末放本表隐藏的 EFF 格）；条件照写全（窗口只是少扫）。
  - 公司小计不引用本列（会成环）：每个账户的数先放隐藏帮助区（40 行，第 a 行＝账户 a），公司行＝SUMIF(帮助 组号, c, 帮助列)。
  - v3.1 录入模式（设计 §15.2）：第 2 行起每个公式都套 layout.pause()（基础资料①「录入模式」＝是 → 全空、不扫 _流）；
    第 1 行标题录入模式时换成 layout.PAUSE_TITLE（红字）。
  - v3.1 WPS 存盘：不写往下撑的区域（$X$8:X8 这种）——第几个用 ROW()-常数；②③ 的序号用隐藏链列 CN（上一格＋本行是不是账户）。
  - 资金日报打印（复核 compat-4）：不缩成一页宽（68 列会缩到 14%），固定 48%、重复 A:B 列、打印区域 ①、②③ 两块。
私有帮助格（不进 layout.NAMES，只本表用）：
  资金日报 BP/BQ（① 每行账户号、窗内笔数）、BR15～BR17（日前窗末、月前窗末、填的截止）、BT..CI（②③ 每个账户）、CK..CN（②③ 每行键、序号链）；
  月度汇报表 R..V（一 每月窗口、月码、有数、显示；二 每个账户的号、笔数）；
  公司汇总 V..Z（组号、本年收入/支出/内部转入/转出）、AB..AE（每月窗口、月码、显示）、AF..AQ（每个账户 12 个月的月末余额）、AS..AU（三 每行键），
  第 37、71 行「未列示」（合计里没落到哪家公司的：账户的公司没认出）。"""
import re
from openpyxl.styles import Font
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from layout import (
    SH_DAY, SH_MON, SH_CO, N_ACC, N_CO, TWO_ROWS, K_DAY, K_YM, ATTR_NO, ATTR_NOINT_HI, CO_METRICS, ACC_TYPES, ACC_STATUS,
    DAY_SEL, DAY_EFF, DAY_KPI, DAY1, DAY2, DAY3, MON_SEL, MON_EFF, MON1, MON2, CO_SEL, CO_EFF, CO1, CO2, CO3, CO_HELP,
    FMT_DATE_HDR, DV_FIXED, pause, PAUSE_TITLE, dsel, dsel_bad, yr_eff, yr_bad, ym_eff, ym_of, raw, win_first, win_last, win_before, win,
    k_aypm, sum_year, sum_before_year)
from common import (
    YH, F_HDR, F_TXT, F_TXTB, F_AUTO, F_AUTOB, F_HELP, F_RED, F_KPI_L, F_KPI_V, F_NOTE, FILL_KPI, FILL_SUB, FILL_AUTO,
    FILL_NONE, BD, AC, ACW, AL, ALW, AR, ACC, DATE, MONTH, INT, PCT, put, title, header, section, widths, hide, selector,
    home_link, cf_flags, cf_neg_red, cf_bg, fill)
from openpyxl.styles import Alignment

I9 = ATTR_NO['内部转账']
A_IN, A_OUT, A_CAP, A_FLOW, A_TBD = (ATTR_NO[k] for k in ('收入', '支出', '投资建设', '往来', '待分类'))
LAST_YM = '(YEAR(P_最后日期)*100+MONTH(P_最后日期))'
OPEN_YM = '(YEAR(P_开账日-1)*100+MONTH(P_开账日-1))'          # 开账日前一天所在的月（月末余额＝期初）；更早的月份不显示
ASHRINK = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)
ACSHRINK = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
F_SUB = Font(name=YH, sz=10, bold=True, color='FF1F3864')


def R(f):
    """录入模式门控（设计 §15.2）：最外层 =IF(P_录入模式=1,"",…)"""
    return pause(f)


def nth(r0):
    """第 r0 行起的第几个（v3.1：不写「第 r0 行到本行」那种往下撑的区域——WPS 存盘会拆）"""
    return f'(ROW()-{r0 - 1})'


def ab(c):
    """'BR1' → '$BR$1'"""
    m = re.fullmatch(r'([A-Z]+)(\d+)', c)
    return f'${m.group(1)}${m.group(2)}'


def d_txt(x):
    return f'YEAR({x})&"/"&MONTH({x})&"/"&DAY({x})'


def ym_txt(x):
    return f'INT({x}/100)&"年"&MOD({x},100)&"月"'


def pause_title(ws, f):
    """第 1 行标题：录入模式时换成 layout.PAUSE_TITLE（红字，条件格式），平时 f。第 1 行不归 pause() 管（make 从第 2 行查）"""
    ws['A1'] = f'=IF(P_录入模式=1,"{PAUSE_TITLE}",{f[1:] if f.startswith("=") else f})'
    ws.conditional_formatting.add('A1', FormulaRule(formula=['P_录入模式=1'], font=Font(name=YH, sz=15, bold=True, color='FFC00000'),
                                                    fill=cf_bg('FFFFD7D7')))


def group_name():
    return 'IF(P_集团名称="","集团",P_集团名称)'


def money(ws, coord, f, bold=False, fill_=None, font=None):
    return put(ws, coord, f, font or (F_AUTOB if bold else F_AUTO), fill_ or FILL_AUTO, fmt=ACC, align=AR)


def txt(ws, coord, f, bold=False, fill_=None, align=None, font=None):
    return put(ws, coord, f, font or (F_AUTOB if bold else F_AUTO), fill_ or FILL_AUTO, align=align or AL)


def helper(ws, coord, f):
    c = ws[coord]
    c.value = f
    c.font = F_HELP
    return c


def merge_put(ws, rng, f, font, fill_, fmt=None, align=AC):
    ws.merge_cells(rng)
    c1 = rng.split(':')[0]
    put(ws, c1, f, font, fill_, fmt=fmt, align=align)
    m = re.fullmatch(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', rng)
    for ci in range(CI(m.group(1)), CI(m.group(3)) + 1):
        for r in range(int(m.group(2)), int(m.group(4)) + 1):
            ws.cell(row=r, column=ci).border = BD


def by_date(sel, eff, empty_note=''):
    """日期选择格旁边：按 … 算 / ※ 看不懂"""
    return R(f'IF({dsel_bad(sel)},"※ 看不懂，按 "&{d_txt(eff)}&" 算","按 "&{d_txt(eff)}&" 算"&IF({raw(sel)},"","{empty_note}"))')


def by_year(sel, eff):
    return R(f'IF({yr_bad(sel)},"※ 看不懂，按 "&{eff}&" 年算",IF({eff}<P_开账年,"※ 本本从 "&P_开账年&" 年开始，"&{eff}&" 年的看上一本",'
             f'"按 "&{eff}&" 年算"&IF({raw(sel)},"","（管理年度）")))')


def by_month(sel, year, eff, empty_note=''):
    bad = f'AND({raw(sel)},{ym_eff(sel, year, 0)}=0)'
    return R(f'IF({bad},"※ 看不懂，按 "&{ym_txt(eff)}&" 算","按 "&{ym_txt(eff)}&" 算"&IF({raw(sel)},"","{empty_note}"))')


def two_level_cf(ws, rng, key_cell, a_cell):
    """两级清单：公司（段）行粗体浅蓝底"""
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'AND({key_cell}<>0,{a_cell}=0)'],
                                                   font=Font(name=YH, bold=True, color='FF1F3864'), fill=cf_bg('FFDDEBF7')))


def two_level_name(key, a, c, sub=' 小计'):
    return (f'IF({key}=0,"",IF({a}>0,"　"&INDEX(账户_名称,{a}),IF({c}=98,"公司没认出的账户（去基础资料③改「公司」）",'
            f'IF({c}=99,"暂停 / 销户的账户",INDEX(公司_名称,{c})&"{sub}"))))')


def two_level_val(key, a, c, hcol, h1, h2, gcol):
    """两级清单一格：账户行＝帮助区第 a 行；公司/段行＝SUMIF(帮助 组号, c, 帮助列)"""
    rg = f'{hcol}${h1}:{hcol}${h2}'
    return f'IF({key}=0,"",IF({a}>0,INDEX({rg},{a}),SUMIF(${gcol}${h1}:${gcol}${h2},{c},{rg})))'


def group_no(a):
    """账户 a 在两级清单里属于哪一组：0 无效；99 暂停/销户；公司号；98 公司没认出（跟 _表 两级候选键一致）"""
    return (f'IF({a}=0,0,IF(INDEX(账户_状态号,{a})>1,99,IF(INDEX(账户_公司号,{a})>0,INDEX(账户_公司号,{a}),98)))')


# ═════════════════════════════ 资金日报 ═════════════════════════════
DAY_PRIV = dict(日前窗末='BR15', 月前窗末='BR16', 止填='BR17')
D_ACC, D_CNT = 'BP', 'BQ'                                   # ① 每行：账户号、窗内笔数
D_H1 = DAY2['首行']                                          # ②③ 帮助区：第 D_H1+a-1 行＝账户 a
D_H2 = D_H1 + N_ACC - 1
D_HC = dict(号='BT', 组='BU', 日笔数='BV', 昨日='BW', 日收='BX', 日支='BY', 日内入='BZ', 日内出='CA', 日余='CB',
            月笔数='CC', 月初='CD', 月收='CE', 月支='CF', 月内入='CG', 月内出='CH', 月末='CI')
D_ROW = dict(键='CK', 号='CL', 组='CM', 序='CN')              # ②③ 每行（两级清单）；序＝到本行为止有几个账户行（链）


def build_day(ws):
    E = {k: ab(v) for k, v in {**DAY_EFF, **DAY_PRIV}.items()}
    last_col = DAY1['期末']
    widths(ws, {'A': 26, 'B': 24})
    c0 = CI(DAY1['首列'])
    for j in range(62):
        ws.column_dimensions[CL(c0 + j)].width = 13
    widths(ws, {'K': 24, DAY1['收款']: 15, DAY1['付款']: 15, DAY1['期末']: 16})
    tip = ('三张小表，全部从【资金台帐】自动算，不用填。① 一段日期（最多 31 天）每个账户每天收了多少、付了多少——账户口径，'
           '收款、付款含内部转入/转出，跟银行流水对得上；② 某一天各账户：昨日余额、本日收入/支出（不含内部转账）、内部转入/转出、本日余额；'
           '③ 某个月各账户：月初、本月收入/支出（不含内部）、内部转入/转出、月末。余额都按日期算（不看台帐上下顺序）。'
           '黄格可以改：① 起始空着＝最后一笔流水那个月的 1 号，截止空着＝起始后 31 天内到最后一笔；② 日期空着＝① 的截止日；'
           '③ 月份空着＝② 那天所在的月（填 1～12＝② 那天所在年份的那个月）。暂停的账户放在 ②③ 最后。'
           '打印：① 每页约 7 天（选的天数以外的页是空的，打印时选页码），②③ 另起一页。')
    title(ws, '=""', 'Q', tip, h2=58)
    pause_title(ws, f'{group_name()}&" 资金日报（"&{d_txt(E["起"])}&"～"&IF({E["止填"]}<{E["起"]},"※ 截止早于起始",{d_txt(E["止"])})&"）"')
    # —— 隐藏的生效值 ——
    lab = dict(起='起始', 止='截止（生效）', 天数='天数', 日='② 日期', 月码='③ 月码', 月初='③ 月初', 月末='③ 月末', 窗首='① 窗首',
               窗末='① 窗末', 前窗末='① 期初窗末', 日窗首='② 窗首', 日窗末='② 窗末', 月窗首='③ 窗首', 月窗末='③ 窗末',
               日前窗末='② 昨日窗末', 月前窗末='③ 月初窗末', 止填='截止（填的或默认）')
    S = {k: ab(v) for k, v in DAY_SEL.items()}
    f = {
        '起': dsel(S['起始'], 'DATE(YEAR(P_最后日期),MONTH(P_最后日期),1)'),
        '止填': dsel(S['截止'], f'MIN({E["起"]}+30,MAX({E["起"]},P_最后日期))'),
        '止': f'IF({E["止填"]}<{E["起"]},{E["起"]}-1,MIN({E["止填"]},{E["起"]}+30))',
        '天数': f'{E["止"]}-{E["起"]}+1',
        '日': dsel(S['日期'], E['止']),
        '月码': ym_eff(S['月份'], f'YEAR({E["日"]})', f'YEAR({E["日"]})*100+MONTH({E["日"]})'),
        '月初': f'DATE(INT({E["月码"]}/100),MOD({E["月码"]},100),1)',
        '月末': f'EOMONTH({E["月初"]},0)',
        '窗首': win_first(ym_of(E['起'])),
        '窗末': win_last(ym_of(E['止']), E['窗首']),
        '前窗末': win_last(ym_of(f'({E["起"]}-1)')),
        '日窗首': win_first(ym_of(E['日'])),
        '日窗末': win_last(ym_of(E['日']), E['日窗首']),
        '月窗首': win_first(E['月码']),
        '月窗末': win_last(E['月码'], E['月窗首']),
        '日前窗末': win_last(ym_of(f'({E["日"]}-1)')),
        '月前窗末': win_before(E['月码']),
    }
    for k, v in f.items():
        cell = {**DAY_EFF, **DAY_PRIV}[k]
        helper(ws, cell, R(v))
        helper(ws, f'BS{cell[2:]}', lab[k])
    for k in ('起', '止', '日', '月初', '月末', '止填'):
        ws[{**DAY_EFF, **DAY_PRIV}[k]].number_format = 'yyyy/mm/dd'
    # —— 第 3 行：选择格 ——
    selector(ws, 'A3', '起始日期', DAY_SEL['起始'], kind='date', prompt='空着＝最后一笔流水那个月的 1 号')
    txt(ws, 'C3', by_date(S['起始'], E['起'], '（默认）'), align=ASHRINK, font=F_NOTE, fill_=FILL_NONE)
    selector(ws, 'D3', '截止日期', DAY_SEL['截止'], kind='date', prompt='空着＝起始后 31 天内到最后一笔流水；最多 31 天')
    end_bad = f'AND({raw(S["截止"])},{dsel_bad(S["截止"])})'
    merge_put(ws, 'F3:G3', R(f'IF({end_bad},"※ 看不懂，按 "&{d_txt(E["止填"])}&" 算",IF({E["止填"]}<{E["起"]},"※ 截止早于起始",'
                             f'"按 "&{d_txt(E["止"])}&" 算"&IF({raw(S["截止"])},"","（默认）")))'),
              F_NOTE, FILL_NONE, align=ASHRINK)
    merge_put(ws, 'H3:I3', R(f'IF({E["止填"]}<{E["起"]},"※ 截止早于起始，请改日期",IF({E["止填"]}-{E["起"]}+1>31,'
                             f'"※ 超过 31 天：只算到 "&{d_txt(E["止"])}&"，请分段查","共 "&{E["天数"]}&" 天"))'),
              F_RED, FILL_NONE, align=ACSHRINK)
    cf_flags(ws, 'H3', 'H3')
    home_link(ws, 'K3')
    ws.row_dimensions[3].height = 22
    # —— 第 4、5 行：四个大数 ——
    r0, rt = DAY1['首行'], DAY1['合计行']
    merge_put(ws, 'A4:A5', '① 资金日报统计表\n（账户口径：收款、付款含内部转入/转出）', F_KPI_L, fill('FFD9E1F2'), align=ACW)
    for (lc, vc, lbl, src) in (('B4:C4', 'B5:C5', '期初余额（起始日前）', f'{DAY1["期初"]}{rt}'),
                               ('D4:E4', 'D5:E5', '期间收款（含内部转入）', f'{DAY1["收款"]}{rt}'),
                               ('F4:G4', 'F5:G5', '期间付款（含内部转出）', f'{DAY1["付款"]}{rt}'),
                               ('H4:I4', 'H5:I5', '期末余额', f'{DAY1["期末"]}{rt}')):
        merge_put(ws, lc, lbl, F_KPI_L, FILL_KPI, align=ACW)
        merge_put(ws, vc, R(f'{src}'), F_KPI_V, FILL_KPI, fmt=ACC)
    assert DAY_KPI == dict(期初='B5', 收款='D5', 付款='F5', 期末='H5', 提示='H3')
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 26
    # —— 第 6、7 行表头 ——
    hr, lr = DAY1['日期行'], DAY1['标签行']
    hfill = fill('FF30561E')
    merge_put(ws, f'A{hr}:A{lr}', '账户', F_HDR, hfill, align=ACW)
    merge_put(ws, f'B{hr}:B{lr}', '期初余额\n（起始日前）', F_HDR, hfill, align=ACW)
    for j in range(31):
        cs, cp = CL(c0 + 2 * j), CL(c0 + 2 * j + 1)
        merge_put(ws, f'{cs}{hr}:{cp}{hr}', R(f'IF({j + 1}<={E["天数"]},{E["起"]}+{j},"")'), F_HDR, hfill, fmt=FMT_DATE_HDR)
        put(ws, f'{cs}{lr}', R(f'IF({cs}${hr}="","","收款")'), F_HDR, hfill, align=AC)
        put(ws, f'{cp}{lr}', R(f'IF({cs}${hr}="","","付款")'), F_HDR, hfill, align=AC)
    for col, t in ((DAY1['收款'], '期间收款'), (DAY1['付款'], '期间付款'), (DAY1['期末'], '期末余额')):
        merge_put(ws, f'{col}{hr}:{col}{lr}', t, F_HDR, hfill, align=ACW)
    ws.row_dimensions[hr].height = 22
    ws.row_dimensions[lr].height = 20
    # —— 第 8～47 行：每个账户 ——
    wk = win('流_账日键', E['窗首'], E['窗末'])
    wsr, wzc = win('流_收', E['窗首'], E['窗末']), win('流_支', E['窗首'], E['窗末'])
    wn0, wk0 = win('流_净', 1, E['前窗末']), win('流_账日键', 1, E['前窗末'])
    for r in range(r0, rt):
        k = nth(r0)
        a = f'${D_ACC}{r}'
        helper(ws, f'{D_ACC}{r}', R(f'IF(INDEX(账户_有效,{k})=1,{k},0)'))
        helper(ws, f'{D_CNT}{r}', R(f'IF(OR({a}=0,{E["天数"]}<1),0,COUNTIFS({wk},">="&({a}*{K_DAY}+{E["起"]}),'
                                    f'{wk},"<="&({a}*{K_DAY}+{E["止"]})))'))
        txt(ws, f'A{r}', R(f'IF({a}=0,"",INDEX(账户_名称,{a}))'), align=ASHRINK, font=F_TXT, fill_=FILL_NONE)
        money(ws, f'{DAY1["期初"]}{r}', R(f'IF({a}=0,"",ROUND(INDEX(账户_期初,{a})+SUMIFS({wn0},{wk0},">="&{a}*{K_DAY},'
                                         f'{wk0},"<"&({a}*{K_DAY}+{E["起"]})),2))'), fill_=FILL_NONE)
        for j in range(31):
            cs, cp = CL(c0 + 2 * j), CL(c0 + 2 * j + 1)
            for col, rng in ((cs, wsr), (cp, wzc)):
                money(ws, f'{col}{r}', R(f'IF(OR({a}=0,{cs}${hr}=""),"",IF(${D_CNT}{r}=0,0,'
                                         f'SUMIF({wk},{a}*{K_DAY}+{cs}${hr},{rng})))'), fill_=FILL_NONE)
        lo, hi = f'${CL(c0)}${lr}', f'${CL(c0 + 61)}${lr}'
        row_rng = f'${CL(c0)}{r}:${CL(c0 + 61)}{r}'
        money(ws, f'{DAY1["收款"]}{r}', R(f'IF({a}=0,"",SUMIF({lo}:{hi},"收款",{row_rng}))'))
        money(ws, f'{DAY1["付款"]}{r}', R(f'IF({a}=0,"",SUMIF({lo}:{hi},"付款",{row_rng}))'))
        money(ws, f'{DAY1["期末"]}{r}', R(f'IF({a}=0,"",ROUND({DAY1["期初"]}{r}+{DAY1["收款"]}{r}-{DAY1["付款"]}{r},2))'),
              bold=True)
    # 合计
    put(ws, f'A{rt}', '合计', F_RED, FILL_KPI, align=AC)
    for col in [DAY1['期初']] + [CL(c0 + j) for j in range(62)] + [DAY1['收款'], DAY1['付款'], DAY1['期末']]:
        if CI(col) >= c0 and CI(col) < c0 + 62:
            ds = CL(c0 + 2 * ((CI(col) - c0) // 2))
            f_ = R(f'IF({ds}${hr}="","",SUM({col}{r0}:{col}{rt - 1}))')
        else:
            f_ = R(f'SUM({col}{r0}:{col}{rt - 1})')
        put(ws, f'{col}{rt}', f_, F_RED, FILL_KPI, fmt=ACC, align=AR)
    cf_neg_red(ws, f'{DAY1["期初"]}{r0}:{DAY1["期初"]}{rt}', f'{DAY1["期初"]}{r0}')
    cf_neg_red(ws, f'{DAY1["期末"]}{r0}:{DAY1["期末"]}{rt}', f'{DAY1["期末"]}{r0}')
    # —— ② ③ ——
    build_day23(ws, E)
    for c in [D_ACC, D_CNT, 'BR', 'BS'] + list(D_HC.values()) + list(D_ROW.values()):
        ws.column_dimensions[c].width = 10
    hide(ws, *([D_ACC, D_CNT, 'BR', 'BS'] + list(D_HC.values()) + list(D_ROW.values())))
    ws.freeze_panes = f'C{r0}'
    print_day(ws)


DAY_PRINT_SCALE = 48      # A4 横向可打宽约 1,065 像素（Excel 按 Calibri 11 算列宽）；②③（A:Q）约 2,127 像素 → ≤50% 才一页宽，留点余量取 48%；① 每页约 7～8 天


def print_day(ws):
    """资金日报打印（复核 compat-4）：v2 套默认「缩成一页宽」，① 68 列挤进一张 A4 横向缩到 14%，看不清。
       改成：A4 横向、固定 48%（不用「缩到一页宽」）、每页重复 A:B 列（账户名、期初）、打印区域两块——① A1:BO48、②③ A50:Q125（另起一页）；
       隐藏的帮助列不在打印区域里。ws._keep_print＝True：make 就不再套默认的一页宽。"""
    ws._keep_print = True
    ps = ws.page_setup
    ps.orientation = 'landscape'
    ps.paperSize = ws.PAPERSIZE_A4
    ps.scale = DAY_PRINT_SCALE
    ps.fitToWidth = ps.fitToHeight = None
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_title_cols = 'A:B'
    ws.print_area = [f'A1:{DAY1["期末"]}{DAY1["合计行"]}', f'A{DAY2["标题行"]}:{DAY3["列"].split(":")[1]}{DAY2["合计行"]}']
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.75
    ws.page_margins.header = ws.page_margins.footer = 0.3


def build_day23(ws, E):
    H = D_HC
    h = lambda k, r: f'{H[k]}{r}'
    S = {k: ab(v) for k, v in DAY_SEL.items()}
    # 帮助区：第 D_H1+a-1 行＝账户 a
    wdk = win('流_账日键', E['日窗首'], E['日窗末'])
    wd = lambda nm: win(nm, E['日窗首'], E['日窗末'])
    wmk = win('流_账月键', E['月窗首'], E['月窗末'])
    wm = lambda nm: win(nm, E['月窗首'], E['月窗末'])
    wy = win('流_账日键', 1, E['日前窗末'])
    wyn = win('流_净', 1, E['日前窗末'])
    wb = win('流_账月键', 1, E['月前窗末'])
    wbn = win('流_净', 1, E['月前窗末'])
    for k, t in (('号', '账户号'), ('组', '两级组'), ('日笔数', '②当天笔数'), ('昨日', '②昨日余额'), ('日收', '②本日收入'),
                 ('日支', '②本日支出'), ('日内入', '②内部转入'), ('日内出', '②内部转出'), ('日余', '②本日余额'),
                 ('月笔数', '③当月笔数'), ('月初', '③月初余额'), ('月收', '③本月收入'), ('月支', '③本月支出'),
                 ('月内入', '③内部转入'), ('月内出', '③内部转出'), ('月末', '③月末余额')):
        helper(ws, f'{H[k]}{D_H1 - 1}', t)
    for r in range(D_H1, D_H2 + 1):
        a = f'${H["号"]}{r}'
        kk = nth(D_H1)
        helper(ws, h('号', r), R(f'IF(INDEX(账户_有效,{kk})=1,{kk},0)'))
        helper(ws, h('组', r), R(group_no(a)))
        dkey = f'{a}*{K_DAY}+{E["日"]}'
        helper(ws, h('日笔数', r), R(f'IF({a}=0,0,COUNTIF({wdk},{dkey}))'))
        helper(ws, h('昨日', r), R(f'IF({a}=0,0,ROUND(INDEX(账户_期初,{a})+SUMIFS({wyn},{wy},">="&{a}*{K_DAY},{wy},"<"&({dkey})),2))'))
        n = f'${H["日笔数"]}{r}'
        for k, val, cond in (('日收', '流_收', f'"<>{I9}"'), ('日支', '流_支', f'"<>{I9}"'), ('日内入', '流_收', I9), ('日内出', '流_支', I9)):
            helper(ws, h(k, r), R(f'IF({n}=0,0,SUMIFS({wd(val)},{wdk},{dkey},{wd("流_属性号")},{cond}))'))
        helper(ws, h('日余', r), R(f'ROUND({h("昨日", r)}+{h("日收", r)}-{h("日支", r)}+{h("日内入", r)}-{h("日内出", r)},2)'))
        mkey = f'{a}*{K_YM}+{E["月码"]}'
        helper(ws, h('月笔数', r), R(f'IF({a}=0,0,COUNTIF({wmk},{mkey}))'))
        helper(ws, h('月初', r), R(f'IF({a}=0,0,ROUND(INDEX(账户_期初,{a})+SUMIFS({wbn},{wb},">="&{a}*{K_YM},{wb},"<"&({mkey})),2))'))
        n = f'${H["月笔数"]}{r}'
        for k, val, cond in (('月收', '流_收', f'"<>{I9}"'), ('月支', '流_支', f'"<>{I9}"'), ('月内入', '流_收', I9), ('月内出', '流_支', I9)):
            helper(ws, h(k, r), R(f'IF({n}=0,0,SUMIFS({wm(val)},{wmk},{mkey},{wm("流_属性号")},{cond}))'))
        helper(ws, h('月末', r), R(f'ROUND({h("月初", r)}+{h("月收", r)}-{h("月支", r)}+{h("月内入", r)}-{h("月内出", r)},2)'))
    # 标题、选择格、表头
    t2, s2, hd, f1, tot = DAY2['标题行'], DAY2['选择行'], DAY2['表头行'], DAY2['首行'], DAY2['合计行']
    assert (t2, s2, hd, f1, tot) == (DAY3['标题行'], DAY3['选择行'], DAY3['表头行'], DAY3['首行'], DAY3['合计行'])
    section(ws, t2, 'A', 'H', '② 货币资金日报表（某一天；账户口径：收入、支出不含内部转账，内部转入/转出单列）')
    section(ws, t2, 'J', 'Q', '③ 多账户资金汇总·月度（账户口径：月末＝月初＋收入－支出＋内部转入－内部转出）')
    selector(ws, f'A{s2}', '日期', DAY_SEL['日期'], kind='date', prompt='空着＝① 的截止日期')
    merge_put(ws, f'C{s2}:H{s2}', by_date(S['日期'], E['日'], '（空着＝① 的截止日期）'), F_NOTE, FILL_NONE, align=AL)
    selector(ws, f'J{s2}', '月份', DAY_SEL['月份'], kind='month',
             prompt='填 1～12＝② 那天所在年份的那个月；也可以填 202609 或那个月的某一天；空着＝② 那天所在的月',
             err='填 1～12（② 那天所在年份的那个月）、202609 这样的年月，或那个月的某一天')     # 复核 F10：跟输入提示一个说法
    merge_put(ws, f'L{s2}:Q{s2}', by_month(S['月份'], f'YEAR({E["日"]})', E['月码'], '（空着＝② 那天所在的月）'), F_NOTE, FILL_NONE,
              align=AL)
    ws.row_dimensions[s2].height = 22
    hdr2 = ['序号', '账户', '昨日余额', '本日收入', '本日支出', '内部转入', '内部转出', '本日余额']
    hdr3 = ['序号', '账户', '月初余额', '本月收入', '本月支出', '内部转入', '内部转出', '月末余额']
    header(ws, hd, list(zip([CL(i) for i in range(1, 9)], hdr2)) + list(zip([CL(i) for i in range(10, 18)], hdr3)), '报表')
    # 每行（两级清单）
    K = D_ROW
    for k, t in (('键', '两级键'), ('号', '账户号'), ('组', '公司号/段'), ('序', '序号链')):
        helper(ws, f'{K[k]}{hd}', t)
    for r in range(f1, tot):
        key, a, c = f'${K["键"]}{r}', f'${K["号"]}{r}', f'${K["组"]}{r}'
        helper(ws, f'{K["键"]}{r}', R(f'IFERROR(SMALL(账两级_候选键,{nth(f1)}),0)'))
        helper(ws, f'{K["号"]}{r}', R(f'MOD({K["键"]}{r},100)'))
        helper(ws, f'{K["组"]}{r}', R(f'INT({K["键"]}{r}/100)'))
        helper(ws, f'{K["序"]}{r}', R(f'N({K["序"]}{r - 1})+IF({a}>0,1,0)'))      # 第 f1 行的上一格是表头文字，N()＝0
        seq = R(f'IF({a}>0,${K["序"]}{r},"")')
        nm = R(two_level_name(key, a, c))
        for (cseq, cname, cols, keys) in (('A', 'B', 'CDEFGH', ('昨日', '日收', '日支', '日内入', '日内出', '日余')),
                                          ('J', 'K', 'LMNOPQ', ('月初', '月收', '月支', '月内入', '月内出', '月末'))):
            put(ws, f'{cseq}{r}', seq, F_AUTO, FILL_NONE, fmt=INT, align=AC)
            put(ws, f'{cname}{r}', nm, F_TXT, FILL_NONE, align=ASHRINK)
            for col, hk in zip(cols, keys):
                money(ws, f'{col}{r}', R(two_level_val(key, a, c, H[hk], D_H1, D_H2, H['组'])), fill_=FILL_NONE,
                      bold=hk in ('日余', '月末'))
    two_level_cf(ws, f'A{f1}:H{tot - 1}', f'${K["键"]}{f1}', f'${K["号"]}{f1}')
    two_level_cf(ws, f'J{f1}:Q{tot - 1}', f'${K["键"]}{f1}', f'${K["号"]}{f1}')
    for col in ('H', 'Q', 'C', 'L'):
        cf_neg_red(ws, f'{col}{f1}:{col}{tot}', f'{col}{f1}')
    # 合计（全部账户：帮助区整列）
    for cseq, cols, keys in (('A', 'CDEFGH', ('昨日', '日收', '日支', '日内入', '日内出', '日余')),
                             ('J', 'LMNOPQ', ('月初', '月收', '月支', '月内入', '月内出', '月末'))):
        merge_put(ws, f'{cseq}{tot}:{CL(CI(cseq) + 1)}{tot}', '合计（全部账户）', F_RED, FILL_KPI)
        for col, hk in zip(cols, keys):
            put(ws, f'{col}{tot}', R(f'SUM({H[hk]}${D_H1}:{H[hk]}${D_H2})'), F_RED, FILL_KPI, fmt=ACC, align=AR)


# ═════════════════════════════ 月度汇报表 ═════════════════════════════
MON_PRIV = dict(窗首='R', 窗末='S', 月码='T', 有数='U', 显示='V')     # 一：第 6..17 行每月
MON2_PRIV = dict(号='R', 笔数='S')                                 # 二：第 23..62 行每个账户


def build_mon(ws):
    E = {k: ab(v) for k, v in MON_EFF.items()}
    S = {k: ab(v) for k, v in MON_SEL.items()}
    C1, C2 = MON1['列'], MON2['列']
    widths(ws, dict(A=24, B=15, C=15, D=15, E=14, F=14, G=14, H=14, I=14, J=14, K=15, L=16, M=40))
    tip = ('一、每个月的经营收支和月末资金总额。经营口径：经营收入、经营支出只算收费项属性是「收入」「支出」的（收入类记在支出栏的算冲减）；'
           '投资建设（正数＝花出去）、往来、待分类、内部转账单列，不算经营；资金净增加＝本月全部收－支；月末资金＝全部账户截至月底的余额。'
           '二、选一个月看各账户（账户口径：收入、支出不含内部转账，内部转入、转出单列；月末＝月初＋收入－支出＋内部转入－内部转出，'
           '跟子表的月份表、资金日报 ③ 一样）。年份空着＝管理年度；最后一笔流水以后的月份不显示。全部自动算，不用填。')
    title(ws, '=""', 'M', tip, h2=58)
    pause_title(ws, f'{group_name()}&" "&{E["年份"]}&" 年月度汇报表"')
    r1, rt = MON1['首行'], MON1['合计行']
    P = MON_PRIV
    for k in MON_PRIV:
        helper(ws, f'{P[k]}{MON1["表头行"]}', f'一·{k}')
    # 生效值
    dflt = (f'IF({E["年份"]}=YEAR(P_最后日期),{LAST_YM},{E["年份"]}*100+IF(MAX(${P["有数"]}${r1}:${P["有数"]}${rt - 1})>0,'
            f'MAX(${P["有数"]}${r1}:${P["有数"]}${rt - 1}),12))')
    f = dict(年份=yr_eff(S['年份']), 月码=ym_eff(S['月份'], E['年份'], dflt),
             月初=f'DATE(INT({E["月码"]}/100),MOD({E["月码"]},100),1)', 月末=f'EOMONTH({E["月初"]},0)',
             窗首=win_first(f'{E["年份"]}*100+1'), 窗末=win_last(f'{E["年份"]}*100+12', E['窗首']),
             月窗首=win_first(E['月码']), 月窗末=win_last(E['月码'], E['月窗首']), 前窗末=win_before(E['月码']))
    for k, v in f.items():
        helper(ws, MON_EFF[k], R(v))
        helper(ws, f'O{MON_EFF[k][1:]}', k)
    for k in ('月初', '月末'):
        ws[MON_EFF[k]].number_format = 'yyyy/mm/dd'
    # 第 3 行
    selector(ws, 'A3', '年份', MON_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    merge_put(ws, 'C3:D3', by_year(S['年份'], E['年份']), F_NOTE, FILL_NONE, align=AL)
    merge_put(ws, 'E3:F3', '本年经营净额', F_KPI_L, FILL_KPI, align=ACW)
    put(ws, 'G3', R(f'{C1["经营净额"]}{rt}'), F_KPI_V, FILL_KPI, fmt=ACC, align=AC)
    merge_put(ws, 'H3:I3', '年末资金总额', F_KPI_L, FILL_KPI, align=ACW)
    merge_put(ws, 'J3:K3', R(f'{C1["月末资金"]}{rt}'), F_KPI_V, FILL_KPI, fmt=ACC)
    home_link(ws, 'M3')
    ws.row_dimensions[3].height = 24
    section(ws, 4, 'A', 'M', '一、按月汇总（经营口径；月末资金是全部账户的余额）')
    hdr = dict(月份='月份', 经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设\n（正数＝花出去）',
               往来流入='往来流入', 往来流出='往来流出', 待分类='待分类净额\n（请定性质）', 内部转入='内部转入', 内部转出='内部转出',
               资金净增加='资金净增加', 月末资金='月末资金总额', 提示='提示')
    header(ws, MON1['表头行'], [(C1[k], hdr[k]) for k in C1], '报表', height=34)
    Y = E['年份']
    for r in range(r1, rt):
        m = nth(r1)
        ym = f'${P["月码"]}{r}'
        helper(ws, f'{P["窗首"]}{r}', R(win_first(f'{Y}*100+{m}')))
        helper(ws, f'{P["窗末"]}{r}', R(win_last(f'{Y}*100+{m}', f'{P["窗首"]}{r}')))
        helper(ws, f'{P["月码"]}{r}', R(f'{Y}*100+{m}'))
        w = lambda nm: win(nm, f'${P["窗首"]}{r}', f'${P["窗末"]}{r}')
        helper(ws, f'{P["有数"]}{r}', R(f'IF(COUNTIF({w("流_月码")},{ym})>0,{m},0)'))
        helper(ws, f'{P["显示"]}{r}', R(f'IF(AND({ym}>={OPEN_YM},OR({ym}<={LAST_YM},{P["有数"]}{r}>0)),1,0)'))
        sh = f'${P["显示"]}{r}=0'
        sp = lambda val, p: f'SUMIFS({w(val)},{w("流_月码")},{ym},{w("流_属性号")},{p})'
        vals = dict(经营收入=sp('流_净', A_IN), 经营支出=f'-{sp("流_净", A_OUT)}',
                    经营净额=f'{C1["经营收入"]}{r}-{C1["经营支出"]}{r}', 投资建设=f'-{sp("流_净", A_CAP)}',
                    往来流入=sp('流_收', A_FLOW), 往来流出=sp('流_支', A_FLOW), 待分类=sp('流_净', A_TBD),
                    内部转入=sp('流_收', I9), 内部转出=sp('流_支', I9), 资金净增加=f'SUMIF({w("流_月码")},{ym},{w("流_净")})')
        put(ws, f'{C1["月份"]}{r}', R(f'DATE({Y},{m},1)'), F_TXTB, FILL_NONE, fmt=MONTH, align=AC)
        for k, v in vals.items():
            money(ws, f'{C1[k]}{r}', R(f'IF({sh},"",{v})'), fill_=FILL_NONE, bold=k in ('经营净额', '资金净增加'))
        wl = lambda nm: win(nm, 1, f'${P["窗末"]}{r}')
        money(ws, f'{C1["月末资金"]}{r}', R(f'IF({sh},"",ROUND(P_期初合计+SUMIF({wl("流_月码")},"<="&{ym},{wl("流_净")}),2))'),
              bold=True, fill_=FILL_NONE)
        txt(ws, f'{C1["提示"]}{r}', R(f'IF({sh},"",{_int_tip(C1, r)})'), font=F_TXT, fill_=FILL_NONE)
    # 全年
    put(ws, f'{C1["月份"]}{rt}', R(f'{Y}&"年全年"'), F_RED, FILL_KPI, align=AC)
    for k in ('经营收入', '经营支出', '经营净额', '投资建设', '往来流入', '往来流出', '待分类', '内部转入', '内部转出', '资金净增加'):
        put(ws, f'{C1[k]}{rt}', R(f'SUM({C1[k]}{r1}:{C1[k]}{rt - 1})'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    wl = lambda nm: win(nm, 1, E['窗末'])
    put(ws, f'{C1["月末资金"]}{rt}', R(f'ROUND(P_期初合计+SUMIF({wl("流_月码")},"<="&({Y}*100+12),{wl("流_净")}),2)'),
        F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{C1["提示"]}{rt}', R(_int_tip(C1, rt, year=True)), F_TXT, FILL_KPI, align=AL)
    cf_flags(ws, f'{C1["提示"]}{r1}:{C1["提示"]}{rt}', f'{C1["提示"]}{r1}')
    for k in ('经营净额', '资金净增加', '待分类'):
        cf_neg_red(ws, f'{C1[k]}{r1}:{C1[k]}{rt}', f'{C1[k]}{r1}')
    build_mon2(ws, E, S)
    hide(ws, 'O', 'P', *MON_PRIV.values())
    ws.freeze_panes = f'B{r1}'


def _int_tip(C1, r, year=False):
    d = f'ROUND({C1["内部转入"]}{r}-{C1["内部转出"]}{r},2)'
    more = '（跨月到账的下个月会对上）' if not year else '（全年应该是 0）'
    return (f'IF({d}=0,"",IF({d}>0,"△ 内部转入比转出多 ","△ 内部转出比转入多 ")&TEXT(ABS({d}),"#,##0.00")&'
            f'"{more}：漏录一头、金额或日期对不上，看核对表清单")')


def build_mon2(ws, E, S):
    C2 = MON2['列']
    P = MON2_PRIV
    t, s, hd, r1, rt = MON2['标题行'], MON2['选择行'], MON2['表头行'], MON2['首行'], MON2['合计行']
    section(ws, t, 'A', 'I', '二、指定月份 · 各账户（账户口径：收入、支出不含内部转账，内部转入/转出单列；月末＝月初＋收入－支出＋内部转入－内部转出）')
    selector(ws, f'A{s}', '月份', MON_SEL['月份'], kind='month',
             prompt='填 1～12＝上面年份的那个月；也可以填 202609 或那个月的某一天；空着＝该年最后有数的月',
             err='填 1～12（上面年份的那个月）、202609 这样的年月，或那个月的某一天')           # 复核 F10
    merge_put(ws, f'C{s}:E{s}', by_month(S['月份'], E['年份'], E['月码'], '（空着＝该年最后有数的月）'), F_NOTE, FILL_NONE, align=AL)
    ws.row_dimensions[s].height = 22
    hdr = dict(账户='账户', 公司='公司', 月初余额='月初余额', 收入='本月收入\n（不含内部）', 支出='本月支出\n（不含内部）', 内部转入='内部转入',
               内部转出='内部转出', 月末余额='月末余额', 占比='占资金比重')
    header(ws, hd, [(C2[k], hdr[k]) for k in C2], '报表', height=34)
    for k in P:
        helper(ws, f'{P[k]}{hd}', f'二·{k}')
    wmk = win('流_账月键', E['月窗首'], E['月窗末'])
    wm = lambda nm: win(nm, E['月窗首'], E['月窗末'])
    wb, wbn = win('流_账月键', 1, E['前窗末']), win('流_净', 1, E['前窗末'])
    for r in range(r1, rt):
        kk = nth(r1)
        a = f'${P["号"]}{r}'
        mkey = f'{a}*{K_YM}+{E["月码"]}'
        helper(ws, f'{P["号"]}{r}', R(f'IF(INDEX(账户_有效,{kk})=1,{kk},0)'))
        helper(ws, f'{P["笔数"]}{r}', R(f'IF({a}=0,0,COUNTIF({wmk},{mkey}))'))
        txt(ws, f'{C2["账户"]}{r}', R(f'IF({a}=0,"",INDEX(账户_名称,{a}))'), font=F_TXT, fill_=FILL_NONE, align=ASHRINK)
        txt(ws, f'{C2["公司"]}{r}', R(f'IF({a}=0,"",IF(INDEX(账户_公司号,{a})>0,INDEX(公司_名称,INDEX(账户_公司号,{a})),"（公司没认出）"))'),
            font=F_TXT, fill_=FILL_NONE, align=ACSHRINK)
        money(ws, f'{C2["月初余额"]}{r}', R(f'IF({a}=0,"",ROUND(INDEX(账户_期初,{a})+SUMIFS({wbn},{wb},">="&{a}*{K_YM},{wb},"<"&({mkey})),2))'),
              fill_=FILL_NONE)
        n = f'${P["笔数"]}{r}'
        for k, val, cond in (('收入', '流_收', f'"<>{I9}"'), ('支出', '流_支', f'"<>{I9}"'), ('内部转入', '流_收', I9),
                             ('内部转出', '流_支', I9)):
            money(ws, f'{C2[k]}{r}', R(f'IF({a}=0,"",IF({n}=0,0,SUMIFS({wm(val)},{wmk},{mkey},{wm("流_属性号")},{cond})))'),
                  fill_=FILL_NONE)
        c = lambda k: f'{C2[k]}{r}'
        money(ws, c('月末余额'), R(f'IF({a}=0,"",ROUND({c("月初余额")}+{c("收入")}-{c("支出")}+{c("内部转入")}-{c("内部转出")},2))'),
              bold=True, fill_=FILL_NONE)
        tot = f'${C2["月末余额"]}${rt}'
        put(ws, c('占比'), R(f'IF(OR({a}=0,ROUND({tot},2)=0),"",{c("月末余额")}/{tot})'), F_AUTO, FILL_NONE, fmt=PCT, align=AR)
    put(ws, f'{C2["账户"]}{rt}', '合计', F_RED, FILL_KPI, align=AC)
    put(ws, f'{C2["公司"]}{rt}', None, F_RED, FILL_KPI)
    for k in ('月初余额', '收入', '支出', '内部转入', '内部转出', '月末余额'):
        put(ws, f'{C2[k]}{rt}', R(f'SUM({C2[k]}{r1}:{C2[k]}{rt - 1})'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'{C2["占比"]}{rt}', R(f'IF(ROUND({C2["月末余额"]}{rt},2)=0,"",SUM({C2["占比"]}{r1}:{C2["占比"]}{rt - 1}))'),
        F_RED, FILL_KPI, fmt=PCT, align=AR)
    for k in ('月初余额', '月末余额'):
        cf_neg_red(ws, f'{C2[k]}{r1}:{C2[k]}{rt}', f'{C2[k]}{r1}')


# ═════════════════════════════ 公司汇总 ═════════════════════════════
CO_PRIV = dict(组='V', 年收='W', 年支='X', 年内入='Y', 年内出='Z')            # 第 2..41 行每个账户
CO_MON = dict(窗首='AB', 窗末='AC', 月码='AD', 显示='AE')                    # 第 2..13 行每个月
CO_ME0 = 'AF'                                                               # AF..AQ：每个账户 1～12 月的月末余额（指标＝月末资金时才算）
CO3_ROW = dict(键='AS', 号='AT', 组='AU')
CO_H1, CO_H2 = 2, 1 + N_ACC
CO_UNL1, CO_UNL2 = CO1['合计行'] + 1, CO2['合计行'] + 1                       # 「未列示」行（私有）
METRIC_PN = {1: [(A_IN, 1)], 2: [(A_OUT, -1)], 3: [(A_IN, 1), (A_OUT, 1)], 4: [(A_CAP, -1)], 5: [(A_FLOW, 1)], 6: [(A_TBD, 1)],
             7: [(I9, 1)]}                                                  # 指标号 → Σ 符号×净（属性 p）
assert CO_METRICS[:7] == ['经营收入', '经营支出', '经营净额', '投资建设', '往来净额', '待分类', '内部净额'] and CO_METRICS[7] == '月末资金'


def build_co(ws):
    E = {k: ab(v) for k, v in CO_EFF.items()}
    S = {k: ab(v) for k, v in CO_SEL.items()}
    Y = E['年份']
    widths(ws, dict(A=26, B=15, C=15, D=15, E=15, F=14, G=14, H=14, I=14, J=15, K=15, L=14, M=14, N=15))
    tip = ('一、各公司一年的账。年初、年末资金按「账户所属公司」算（这家公司名下账户的余额）；经营收入/支出、投资建设、往来、待分类、内部转账按「归属公司」算'
           '（私户替别的公司收付时台帐里填的归属公司，没填＝账户所属公司），所以一家公司的「年末－年初」跟各项之和会差一点，差额列在「代收代付差额」，'
           '全集团合计是 0。二、各公司每个月的一项指标（E3 选；月末资金按账户所属公司，其余按归属公司）。'
           '三、各公司名下每个账户（账户口径：收入、支出不含内部转账；年末＝年初＋收入－支出＋内部转入－内部转出），暂停/销户的账户放最后。'
           '年份空着＝管理年度；最后一笔流水以后的月份不显示。全部自动算，不用填。')
    title(ws, '=""', 'N', tip, h2=58)
    pause_title(ws, f'{group_name()}&" "&{Y}&" 年公司汇总"')
    metrics = '"' + '","'.join(CO_METRICS) + '"'
    sel_m = f'TRIM({S["指标"]}&"")'
    f = dict(年份=yr_eff(S['年份']),
             指标号=f'IF({sel_m}="",{len(CO_METRICS)},IFERROR(MATCH({sel_m},{{{",".join(chr(34) + m + chr(34) for m in CO_METRICS)}}},0),'
                    f'{len(CO_METRICS)}))',
             窗首=win_first(f'{Y}*100+1'), 窗末=win_last(f'{Y}*100+12', E['窗首']), 前窗末=win_before(f'{Y}*100+1'))
    for k, v in f.items():
        helper(ws, CO_EFF[k], R(v))
        helper(ws, f'O{CO_EFF[k][1:]}', k)
    # 第 3 行
    selector(ws, 'A3', '年份', CO_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    txt(ws, 'C3', by_year(S['年份'], Y), font=F_NOTE, fill_=FILL_NONE, align=ASHRINK)
    selector(ws, 'D3', '二表看', CO_SEL['指标'], dv_formula='"' + DV_FIXED['指标'] + '"',
             prompt='选二表（公司×月）看哪一项；空着＝月末资金')
    mname = f'CHOOSE({E["指标号"]},{metrics})'
    merge_put(ws, 'F3:G3', R(f'IF(AND({sel_m}<>"",ISERROR(MATCH({sel_m},{{{",".join(chr(34) + m + chr(34) for m in CO_METRICS)}}},0))),'
                             f'"※ 看不懂，按 "&{mname}&" 算","按 "&{mname}&" 算"&IF({sel_m}="","（默认）",""))'),
              F_NOTE, FILL_NONE, align=AL)
    home_link(ws, 'I3')
    ws.row_dimensions[3].height = 24
    # —— 帮助区：每个账户 ——
    H = {**CO_HELP, **CO_PRIV}
    for k, c in H.items():
        helper(ws, f'{c}1', k)
    wn0, wk0 = win('流_净', 1, E['前窗末']), win('流_账年属月键', 1, E['前窗末'])
    wy = lambda nm: win(nm, E['窗首'], E['窗末'])
    wky = wy('流_账年属月键')
    for r in range(CO_H1, CO_H2 + 1):
        kk = nth(CO_H1)
        a = f'${H["账户号"]}{r}'
        helper(ws, f'{H["账户号"]}{r}', R(f'IF(INDEX(账户_有效,{kk})=1,{kk},0)'))
        helper(ws, f'{H["公司号"]}{r}', R(f'IF({a}=0,0,INDEX(账户_公司号,{a}))'))
        helper(ws, f'{H["年初"]}{r}', R(f'IF({a}=0,0,ROUND(INDEX(账户_期初,{a})+{sum_before_year(wn0, wk0, a, Y)},2))'))
        helper(ws, f'{H["年末"]}{r}', R(f'IF({a}=0,0,ROUND(${H["年初"]}{r}+{sum_year(wy("流_净"), wky, a, Y)},2))'))
        helper(ws, f'{H["组"]}{r}', R(group_no(a)))
        helper(ws, f'{H["年收"]}{r}', R(f'IF({a}=0,0,{sum_year(wy("流_收"), wky, a, Y, p=1, p2=ATTR_NOINT_HI)})'))
        helper(ws, f'{H["年支"]}{r}', R(f'IF({a}=0,0,{sum_year(wy("流_支"), wky, a, Y, p=1, p2=ATTR_NOINT_HI)})'))
        helper(ws, f'{H["年内入"]}{r}', R(f'IF({a}=0,0,{sum_year(wy("流_收"), wky, a, Y, p=I9)})'))
        helper(ws, f'{H["年内出"]}{r}', R(f'IF({a}=0,0,{sum_year(wy("流_支"), wky, a, Y, p=I9)})'))
    # —— 帮助区：每个月（第 2..13 行）——
    M = CO_MON
    for k, c in M.items():
        helper(ws, f'{c}1', '月·' + k)
    mrow = {m: 1 + m for m in range(1, 13)}
    for m in range(1, 13):
        r = mrow[m]
        mm = nth(2)
        helper(ws, f'{M["窗首"]}{r}', R(win_first(f'{Y}*100+{mm}')))
        helper(ws, f'{M["窗末"]}{r}', R(win_last(f'{Y}*100+{mm}', f'{M["窗首"]}{r}')))
        helper(ws, f'{M["月码"]}{r}', R(f'{Y}*100+{mm}'))
        wmm = win('流_月码', f'{M["窗首"]}{r}', f'{M["窗末"]}{r}')
        helper(ws, f'{M["显示"]}{r}', R(f'IF(AND({M["月码"]}{r}>={OPEN_YM},OR({M["月码"]}{r}<={LAST_YM},COUNTIF({wmm},{M["月码"]}{r})>0)),1,0)'))
    # —— 帮助区：每个账户 12 个月的月末余额（指标＝月末资金才算）——
    for m in range(1, 13):
        col = CL(CI(CO_ME0) + m - 1)
        helper(ws, f'{col}1', f'{m}月末')
        mr = mrow[m]
        wk = win('流_账月键', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
        wn = win('流_净', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
        prev = f'${H["年初"]}' if m == 1 else CL(CI(CO_ME0) + m - 2)
        for r in range(CO_H1, CO_H2 + 1):
            a = f'${H["账户号"]}{r}'
            helper(ws, f'{col}{r}', R(f'IF(OR({E["指标号"]}<>8,{a}=0),0,ROUND({prev}{r}+SUMIF({wk},{a}*{K_YM}+${M["月码"]}${mr},{wn}),2))'))
    hs, hy = f'${H["公司号"]}${CO_H1}:${H["公司号"]}${CO_H2}', lambda k: f'${H[k]}${CO_H1}:${H[k]}${CO_H2}'
    # —— 一、公司年度汇总 ——
    C = CO1['列']
    r1, rt = CO1['首行'], CO1['合计行']
    section(ws, 4, 'A', 'K', '一、公司年度汇总（资金按账户所属公司；经营收支、投资建设、往来、待分类、内部按归属公司）')
    hdr = dict(公司='公司', 年初资金='年初资金', 经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设\n（正数＝花出去）',
               往来净额='往来净额', 待分类='待分类净额', 内部净额='内部转账净额', 年末资金='年末资金', 代收代付差额='代收代付差额\n（合计应为 0）')
    header(ws, CO1['表头行'], [(C[k], hdr[k]) for k in C], '报表', height=34)
    wg = wy('流_归年属月键')
    for r in range(r1, rt):
        cc = nth(r1)
        nm = f'$A{r}'
        sy = lambda p, sg=1: (('-' if sg < 0 else '') + sum_year(wy('流_净'), wg, cc, Y, p=p))
        txt(ws, f'{C["公司"]}{r}', R(f'IF(INDEX(公司_有效,{cc})=1,INDEX(公司_名称,{cc}),"")'), font=F_TXTB, fill_=FILL_NONE)
        vals = dict(年初资金=f'SUMIF({hs},{cc},{hy("年初")})', 经营收入=sy(A_IN), 经营支出=sy(A_OUT, -1),
                    经营净额=f'{C["经营收入"]}{r}-{C["经营支出"]}{r}', 投资建设=sy(A_CAP, -1), 往来净额=sy(A_FLOW), 待分类=sy(A_TBD),
                    内部净额=sy(I9), 年末资金=f'SUMIF({hs},{cc},{hy("年末")})',
                    代收代付差额=(f'ROUND({C["年末资金"]}{r}-{C["年初资金"]}{r}-{C["经营净额"]}{r}+{C["投资建设"]}{r}-{C["往来净额"]}{r}'
                            f'-{C["待分类"]}{r}-{C["内部净额"]}{r},2)'))
        for k, v in vals.items():
            money(ws, f'{C[k]}{r}', R(f'IF({nm}="","",{v})'), fill_=FILL_NONE, bold=k in ('年末资金', '经营净额'))
    # 合计（全集团，所有账户、所有归属）
    put(ws, f'A{rt}', '合计（全集团）', F_RED, FILL_KPI, align=AC)
    wm_ = wy('流_月码')
    gy = lambda p, sg=1: (('-' if sg < 0 else '') +
                          f'SUMIFS({wy("流_净")},{wm_},">="&({Y}*100+1),{wm_},"<="&({Y}*100+12),{wy("流_属性号")},{p})')
    tv = dict(年初资金=f'SUM({hy("年初")})', 经营收入=gy(A_IN), 经营支出=gy(A_OUT, -1), 经营净额=f'{C["经营收入"]}{rt}-{C["经营支出"]}{rt}',
              投资建设=gy(A_CAP, -1), 往来净额=gy(A_FLOW), 待分类=gy(A_TBD), 内部净额=gy(I9), 年末资金=f'SUM({hy("年末")})',
              代收代付差额=(f'ROUND({C["年末资金"]}{rt}-{C["年初资金"]}{rt}-{C["经营净额"]}{rt}+{C["投资建设"]}{rt}-{C["往来净额"]}{rt}'
                      f'-{C["待分类"]}{rt}-{C["内部净额"]}{rt},2)'))
    for k, v in tv.items():
        put(ws, f'{C[k]}{rt}', R(v), F_RED, FILL_KPI, fmt=ACC, align=AR)
    _unlisted(ws, CO_UNL1, r1, rt, [C[k] for k in C if k != '公司'])
    for k in ('经营净额', '年末资金', '年初资金', '代收代付差额'):
        cf_neg_red(ws, f'{C[k]}{r1}:{C[k]}{CO_UNL1}', f'{C[k]}{r1}')
    # —— 二、公司 × 月 ——
    t2, h2, s2, e2 = CO2['标题行'], CO2['表头行'], CO2['首行'], CO2['合计行']
    ws.merge_cells(f'A{t2}:N{t2}')
    put(ws, f'A{t2}', R(f'"二、公司 × 月："&{mname}&IF({E["指标号"]}=8,"（按账户所属公司，月底余额）","（按归属公司，当月发生）")&'
                        f'"　　在 E3 换指标"'), Font(name=YH, sz=11, bold=True, color='FFFFFFFF'), fill('FF30561E'), align=AL)
    header(ws, h2, [('A', '公司')] + [(CL(1 + m), f'{m}月') for m in range(1, 13)] + [('N', '全年')], '报表')
    put(ws, f'N{h2}', R(f'IF({E["指标号"]}=8,"年末","全年")'), Font(name=YH, sz=10, bold=True, color='FFFFFFFF'), fill('FF30561E'),
        align=AC)
    for r in range(s2, e2):
        cc = nth(s2)
        txt(ws, f'A{r}', R(f'IF(INDEX(公司_有效,{cc})=1,INDEX(公司_名称,{cc}),"")'), font=F_TXTB, fill_=FILL_NONE)
        for m in range(1, 13):
            col, mr = CL(1 + m), mrow[m]
            wk = win('流_归年属月键', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
            wn = win('流_净', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
            sp = lambda p: f'SUMIF({wk},{k_aypm(cc, Y, p, m)},{wn})'
            me = CL(CI(CO_ME0) + m - 1)
            v = _metric_if(E['指标号'], sp, f'SUMIF({hs},{cc},{me}${CO_H1}:{me}${CO_H2})')
            money(ws, f'{col}{r}', R(f'IF(OR($A{r}="",${M["显示"]}${mr}=0),"",{v})'), fill_=FILL_NONE)
        money(ws, f'N{r}', R(f'IF($A{r}="","",IF({E["指标号"]}=8,SUMIF({hs},{cc},{hy("年末")}),SUM(B{r}:M{r})))'), bold=True,
              fill_=FILL_NONE)
    put(ws, f'A{e2}', '合计（全集团）', F_RED, FILL_KPI, align=AC)
    for m in range(1, 13):
        col, mr = CL(1 + m), mrow[m]
        wmm = win('流_月码', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
        wn = win('流_净', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
        wp = win('流_属性号', f'${M["窗首"]}${mr}', f'${M["窗末"]}${mr}')
        sp = lambda p: f'SUMIFS({wn},{wmm},${M["月码"]}${mr},{wp},{p})'
        me = CL(CI(CO_ME0) + m - 1)
        v = _metric_if(E['指标号'], sp, f'SUM({me}${CO_H1}:{me}${CO_H2})')
        put(ws, f'{col}{e2}', R(f'IF(${M["显示"]}${mr}=0,"",{v})'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    put(ws, f'N{e2}', R(f'IF({E["指标号"]}=8,SUM({hy("年末")}),SUM(B{e2}:M{e2}))'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    _unlisted(ws, CO_UNL2, s2, e2, [CL(1 + m) for m in range(1, 14)])
    cf_neg_red(ws, f'B{s2}:N{CO_UNL2}', f'B{s2}')
    # —— 三、公司 → 账户余额（两级）——
    C3 = CO3['列']
    t3, h3, f3, e3 = CO3['标题行'], CO3['表头行'], CO3['首行'], CO3['合计行']
    section(ws, t3, 'A', 'I', '三、公司 → 账户余额（账户口径：收入、支出不含内部转账；年末＝年初＋收入－支出＋内部转入－内部转出；公司行是小计）')
    hdr3 = dict(名称='公司 / 账户', 类型='类型', 状态='状态', 年初余额='年初余额', 本年收入='本年收入\n（不含内部）', 本年支出='本年支出\n（不含内部）',
                内部转入='内部转入', 内部转出='内部转出', 年末余额='年末余额')
    header(ws, h3, [(C3[k], hdr3[k]) for k in C3], '报表', height=34)
    K = CO3_ROW
    for k, c in K.items():
        helper(ws, f'{c}{h3}', '三·' + k)
    src = dict(年初余额='年初', 本年收入='年收', 本年支出='年支', 内部转入='年内入', 内部转出='年内出', 年末余额='年末')
    types = '"' + '","'.join(ACC_TYPES) + '"'
    stats = '"' + '","'.join(ACC_STATUS) + '"'
    for r in range(f3, e3):
        key, a, c = f'${K["键"]}{r}', f'${K["号"]}{r}', f'${K["组"]}{r}'
        helper(ws, f'{K["键"]}{r}', R(f'IFERROR(SMALL(账两级_候选键,{nth(f3)}),0)'))
        helper(ws, f'{K["号"]}{r}', R(f'MOD({K["键"]}{r},100)'))
        helper(ws, f'{K["组"]}{r}', R(f'INT({K["键"]}{r}/100)'))
        txt(ws, f'{C3["名称"]}{r}', R(two_level_name(key, a, c, sub='')), font=F_TXT, fill_=FILL_NONE, align=ASHRINK)
        txt(ws, f'{C3["类型"]}{r}', R(f'IF({a}=0,"",IFERROR(CHOOSE(INDEX(账户_类型号,{a}),{types}),""))'), font=F_TXT, fill_=FILL_NONE,
            align=AC)
        txt(ws, f'{C3["状态"]}{r}', R(f'IF({a}=0,"",IFERROR(CHOOSE(INDEX(账户_状态号,{a}),{stats}),"")&'
                                     f'IF(AND(INDEX(账户_状态号,{a})>1,INDEX(账户_停用日,{a})>0),"（"&MONTH(INDEX(账户_停用日,{a}))&"/"&'
                                     f'DAY(INDEX(账户_停用日,{a}))&" 起）",""))'), font=F_TXT, fill_=FILL_NONE, align=ACSHRINK)
        for k, hk in src.items():
            money(ws, f'{C3[k]}{r}', R(two_level_val(key, a, c, H[hk], CO_H1, CO_H2, H['组'])), fill_=FILL_NONE,
                  bold=k == '年末余额')
    two_level_cf(ws, f'A{f3}:I{e3 - 1}', f'${K["键"]}{f3}', f'${K["号"]}{f3}')
    ws.merge_cells(f'A{e3}:C{e3}')
    put(ws, f'A{e3}', '合计（全部账户）', F_RED, FILL_KPI, align=AC)
    for k, hk in src.items():
        put(ws, f'{C3[k]}{e3}', R(f'SUM({hy(hk)})'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    for k in ('年初余额', '年末余额'):
        cf_neg_red(ws, f'{C3[k]}{f3}:{C3[k]}{e3}', f'{C3[k]}{f3}')
    hide(ws, 'O', 'P', *CO_HELP.values(), *CO_PRIV.values(), *CO_MON.values(), *[CL(CI(CO_ME0) + i) for i in range(12)],
         *CO3_ROW.values())
    ws.freeze_panes = f'B{CO1["首行"]}'


def _metric_if(no, sp, me):
    """指标号 → 公式（IF 只算选中的那一种）：1..7 按 METRIC_PN 拼 Σ 符号×SUMIF(属性 p)；8 月末资金＝me"""
    def one(i):
        out = ''
        for p, sg in METRIC_PN[i]:
            out += ('-' if sg < 0 else ('+' if out else '')) + sp(p)
        return out
    f = one(7)
    for i in range(6, 0, -1):
        f = f'IF({no}={i},{one(i)},{f})'
    return f'IF({no}=8,{me},{f})'


def _unlisted(ws, r, r1, rt, cols):
    """「未列示」：合计里没落到上面哪一行的（账户的公司没认出）＝合计－上面各行；全是 0 整行空着"""
    diff = lambda col: f'ROUND(N({col}{rt})-SUM({col}{r1}:{col}{rt - 1}),2)'
    anyv = 'OR(' + ','.join(f'{diff(c)}<>0' for c in cols) + ')'
    put(ws, f'A{r}', R(f'IF({anyv},"↑ 合计里含：没落到哪家公司的（账户的公司没认出，去基础资料③改）","")'), F_NOTE, FILL_NONE,
        align=ALW)
    ws.row_dimensions[r].height = 30
    for c in cols:
        put(ws, f'{c}{r}', R(f'IF({diff(c)}=0,"",{diff(c)})'), F_NOTE, FILL_NONE, fmt=ACC, align=AR)


def build(wb, ctx):
    build_day(wb[SH_DAY])
    build_mon(wb[SH_MON])
    build_co(wb[SH_CO])
