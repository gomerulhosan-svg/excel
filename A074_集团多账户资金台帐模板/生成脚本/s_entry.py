# -*- coding: utf-8 -*-
"""a74split 录入本的轻公式、首页、使用说明（两本各一份）。split.make_entry / make_query 调；一本版不用本模块。

录入本【资金台帐】（录入列 A..G、I..K、P 跟一本版一模一样；列位不挪）：
  · 右边隐藏辅助列 layout.ENTRY_HELP（R..AA）：每行一样、只看本行的录入格和 _表（照一本版 _流 同名列的写法，项推、人推写成单行公式），
    改一格只算这一行；
  · 显示列：H 余额＝期初＋块前累计（_块，按行号每 100 行一块）＋本块到本行本账户的净额（INDEX(整列名称,第几格) 取区域：排序、插删行都按
    现在的行号算，每行公式一样）；L 问题＝只留逐行就能判断的 ※/△（配对、跳期、重复粘贴、日期顺序这些整表的检查在查询本核对表）；
    M 属性、N/O 认出来的收费项/收付人＝看本行辅助列；
  · 第 1 行 F1 红字：表头/列挪了、有行没有自动列（中间插了行或整行清了）、满了；第 3 行汇总条只留轻的。
录入本【首页】（轻）：资金总额、流水笔数、最后一笔、※/△ 行数、各账户当前余额（_块 的块净按列加），一句「录的时候只开这本」。
_块（隐藏）：第 b 块在第 b+1 行；A 本块有没有数；B 起 N_ACC 列块净（SUMIFS 本块 100 行）；再 N_ACC 列块前累计（上一行前＋上一行净）。
"""
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.worksheet.protection import SheetProtection
from layout import (
    SH_HOME, SH_LEDGER, SH_BASE, SH_HELP, H_PAR, H_TAB, HDR, R0, N_LEDGER, N_ACC, LEDGER_COLS, LEDGER_SHOW, LEDGER_LAST,
    LEDGER_LAST_ROW, ATTR_NO, ATTR_UNKNOWN, BASE_END_ROW, BASE_COLS, BASE_BLOCKS, PAR, PAR_ROW, BASE_PAR_VAL,
    WB1_FILE, WB2_FILE, SH_BLK_E, ENTRY_HELP, ENTRY_HELP_NAME, ENTRY_BLK, ENTRY_BLK_ROWS, ENTRY_N_BLK, ENTRY_KPI,
    ENTRY_KPI_NOTE, FETCH_LAST, RGB_ERR_FONT, RGB_ERR_FILL, RGB_WARN_FONT, C_HOME,
    q, raw, name2no, vdate, chain, esc, norm)
from common import (
    YH, F_TITLE, F_TIP, F_HDR, F_TXT, F_TXTB, F_AUTO, F_AUTOB, F_NOTE, F_HELP, F_KPI_L, F_KPI_V, F_SEC, F_LINK, FILL_TITLE,
    FILL_TIP, FILL_AUTO, FILL_KPI, FILL_NONE, FILL_IN, AC, ACW, AL, ALW, AR, MONEY, DATE, put, title, header, widths, link,
    home_link, kpi, cf_bg, cf_flags, cf_neg_red, fill, print_setup)
from s_mirror import item_lookup, item_multi, RI, RQ, RP, _tail, T_REFUND_IN, T_REFUND_OUT, T_NOPARTY

INT9 = ATTR_NO['内部转账']
_ATTR_NAMES = {v: k for k, v in ATTR_NO.items()}
GATE = ('日期', '账户', '收费项', '收付人', '摘要', '收入', '支出', '对方账户', '归属公司')   # 「这一行有东西」看这 9 格（同一本版）
EB = ENTRY_BLK_ROWS
NM = ENTRY_HELP_NAME
BLK_NET = '录本_块净'
BLK_PRE = '录本_块前'
NEAR = N_LEDGER // 10


def numv2(xx):
    """金额 → 数（同一本版 _流 收值/支值）：认不出＝文字 "坏" """
    tt = f'TRIM({xx}&"")'
    ss = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({tt},",",""),"，",""),"¥",""),"￥","")'
    return (f'IF(ISNUMBER({xx}),ROUND({xx},2),IF(OR({tt}="",{tt}="-",{tt}="--",{tt}="—",{tt}="－"),0,'
            f'IFERROR(ROUND(--{ss},2),IF(ISTEXT({xx}),"坏",0))))')


def _refs(r):
    x = {f: f'${c}{r}' for f, c in LEDGER_COLS.items()}
    h = {f: f'${c}{r}' for f, c in ENTRY_HELP.items()}
    return x, h


def help_formulas(r):
    """第 r 行辅助列（只看本行）→ {字段: 公式}"""
    x, h = _refs(r)
    T, U, S, Xo, Y = h['收值'], h['支值'], h['账户号'], h['对方号'], h['项号']
    dirno = f'(2-OR(N({T})>0,N({U})<0))'
    cono = f'INDEX(账户_公司号,{S})'
    internal = f'AND({Xo}>0,{Xo}<>{S},P_内部号>0)'
    memo, party = x['摘要'], x['收付人']
    guess = item_lookup(RI, RQ, 'P_Q规则数', '规_项号', memo, party, dirno, cono, S)
    cond_p = f'ISNUMBER(SEARCH({RP("词")},{memo}))*({RP("位")}<>2)' + _tail(RP, dirno, cono, S)
    unk = f'IF(OR(N({T})>0,N({U})<0),1,2)' if ATTR_UNKNOWN == 0 else str(ATTR_UNKNOWN)
    return {
        '日期值': f'=IFERROR({vdate(x["日期"])},0)',
        '账户号': f'=IFERROR({name2no(x["账户"], "账户")},0)',
        '收值': f'=IFERROR({numv2(x["收入"])},"坏")',
        '支值': f'=IFERROR({numv2(x["支出"])},"坏")',
        '有效': f'=IF(AND({h["日期值"]}>=P_开账日,{h["日期值"]}<=P_日期上限,{S}>0,OR(N({T})<>0,N({U})<>0)),1,0)',
        '净': f'=IF({h["有效"]}=1,ROUND(N({T})-N({U}),2),0)',
        '对方号': f'=IFERROR({name2no(x["对方账户"], "账户")},0)',
        '项号': (f'=IFERROR(IF({raw(x["收费项"])},{name2no(x["收费项"], "项")},IF({internal},P_内部号,'
               f'IF(AND({S}>0,P_项规则数>0,OR({raw(memo)},{raw(party)})),{guess},0))),0)'),
        '人号': (f'=IFERROR(IF({raw(party)},{name2no(party, "人")},IF(AND({S}>0,P_人规则数>0,{raw(memo)}),'
               f'IFERROR(LOOKUP(1,0/({cond_p}),{RP("人号")}),0),0)),0)'),
        '属性号': (f'=IF({Y}>0,IF(INDEX(项_属性号,{Y})>0,INDEX(项_属性号,{Y}),IF(OR(N({T})>0,N({U})<0),1,2)),'
                f'IF(AND({Xo}>0,{Xo}<>{S}),{INT9},{unk}))'),
    }


def show_formulas(r):
    """第 r 行显示列（H、L、M、N、O）→ {字段: 公式}"""
    x, h = _refs(r)
    T, U, S, Xo, Y, Z, K = (h[k] for k in ('收值', '支值', '账户号', '对方号', '项号', '人号', '属性号'))
    b0 = f'INT((ROW()-{R0})/{EB})'                     # 第几块－1
    s = f'{b0}*{EB}+{R0 - HDR + 1}'                    # 本块第一行在整列名称（表头行起）里是第几格
    e = f'ROW()-{HDR - 1}'                             # 本行是第几格
    bal = (f'=IF({h["有效"]}<>1,"",ROUND(INDEX(账户_期初,{S})+INDEX({BLK_PRE},{b0}+1,{S})+'
           f'SUMIFS(INDEX({NM["净"]},{s}):INDEX({NM["净"]},{e}),INDEX({NM["账户号"]},{s}):INDEX({NM["账户号"]},{e}),{S}),2))')
    choose = ','.join(f'"{_ATTR_NAMES.get(i, "")}"' for i in range(1, max(ATTR_NO.values()) + 1))
    tx = lambda f: f'TRIM({x[f]}&"")'
    # ※（只看本行）：同一本版 问本 的顺序（收费项不在④改成 △，放到下面）
    D_pos = f'OR(N({T})>0,N({U})<0)'                  # 一本版的 收>0
    E_pos = f'OR(N({U})>0,N({T})<0)'                  # 一本版的 支>0
    allf = '&'.join(x[f] for f in LEDGER_COLS)
    cat9 = '&'.join(x[f] for f in GATE)
    err = [
        (f'{tx("日期")}=""', '"※ 没填日期"'),
        (f'{h["日期值"]}=0', '"※ 日期看不懂：要写 2026/9/1 这样的日期"'),
        (f'{h["日期值"]}<P_开账日', '"※ 早于开账日：上一年的账记在上一本"'),
        (f'{h["日期值"]}>P_日期上限', '"※ 日期晚于 "&YEAR(P_日期上限)&" 年底：年份写错了？"'),
        (f'NOT({raw(x["账户"])})', '"※ 没选账户"'),
        (f'{S}=0', f'"※ 账户「"&{tx("账户")}&"」不在基础资料③：改了名就把旧名加进别名"'),
        (f'OR(ISTEXT({T}),ISTEXT({U}))', '"※ 金额看不懂：10000+200 要写成 =10000+200"'),
        (f'AND(N({T})=0,N({U})=0)', '"※ 没填金额"'),
        (f'AND(N({T})<>0,N({U})<>0)', '"※ 收入、支出都填了：分成两行记"'),
        (f'AND({K}={INT9},NOT({raw(x["对方账户"])}))', '"※ 内部转账要选转账的另一户；对方不在本册就改成往来款、借款等"'),
        (f'IF({raw(x["对方账户"])},{Xo}=0,FALSE)', f'"※ 转账的另一户「"&{tx("对方账户")}&"」不在③：外单位、个人请填在收付人"'),
        (f'AND({Xo}>0,{Xo}={S})', '"※ 转账的另一户就是本账户"'),
        (f'AND({Xo}>0,{K}<>{INT9})', '"※ 填了转账的另一户，收费项要选内部转账（原性质写摘要）"'),
        (f'IF({raw(x["归属公司"])},{name2no(x["归属公司"], "公司")}=0,FALSE)', f'"※ 归属公司「"&{tx("归属公司")}&"」不在②"'),
    ]
    memo, party = x['摘要'], x['收付人']
    dirno = f'(2-{D_pos})'
    cono = f'INDEX(账户_公司号,{S})'
    guessed = f'AND(NOT({raw(x["收费项"])}),{Xo}=0,{Y}>0)'
    multi = item_multi(RI, RQ, 'P_Q规则数', memo, party, dirno, cono, S, Y)
    iname = f'INDEX(项_名称,{Y})'
    textc = f'OR(ISTEXT({x["日期"]}),AND(ISTEXT({x["收入"]}),{raw(x["收入"])}),AND(ISTEXT({x["支出"]}),{raw(x["支出"])}))'
    me = f'{x["摘要"]}&{x["备注"]}'
    refund = f'OR(ISNUMBER(SEARCH("退",{me})),ISNUMBER(SEARCH("冲",{me})),ISNUMBER(SEARCH("返还",{me})))'
    paused = (f'AND(INDEX(账户_状态号,{S})>1,OR(INDEX(账户_停用日,{S})=0,{h["日期值"]}>INDEX(账户_停用日,{S})))')
    warn = [
        (f'IF({raw(x["收费项"])},{Y}=0,FALSE)', f'"△ 收费项「"&{tx("收费项")}&"」不在④，报表算待分类：从下拉选，或去④加（改了名就加别名）"'),
        (f'AND({Y}>0,{Y}=P_待定号)', '"△ 收费项是「待定」：定了性质请改过来"'),
        (guessed, f'IF({multi}>0,"△ 摘要能对上好几种收费项，按「"&{iname}&"」算了：请在收费项里选",'
                  f'"△ 收费项按摘要认成「"&{iname}&"」：对就在收费项选上")'),
        (f'AND(NOT({raw(x["收费项"])}),{Y}=0)', '"△ 收费项没填，摘要也认不出：请选收费项"'),
        (textc, '"△ 日期或金额是文字，已照认：最好改成真日期、数字"'),
        (f'OR(N({T})<0,N({U})<0)', '"△ 金额是负数，已当另一栏算"'),
        (f'IF(OR(AND({K}=1,{E_pos}),AND({K}=2,{D_pos})),NOT({refund}),FALSE)', f'IF({K}=1,"{T_REFUND_IN}","{T_REFUND_OUT}")'),
        (f'IF({raw(party)},{Z}=0,FALSE)', f'"△ 收付人「"&{tx("收付人")}&"」不在⑤：报表算没填"'),
        (f'AND({Z}=0,OR({K}=1,{K}=2))', f'"{T_NOPARTY}"'),
        (paused, '"△ 账户已暂停/销户，还在记账"'),
    ]
    prob = (f'=IF(ISERROR({allf}),"※ 这一行有错误值（#N/A、#VALUE! 这种）：改成正常的日期、文字、数字",'
            f'IF(TRIM({cat9})="","",{chain(err + warn)}))')
    return {
        '余额': bal,
        '问题': prob,
        '属性': f'=IF({h["有效"]}<>1,"",IFERROR(CHOOSE({K},{choose}),""))',
        '收费项认': f'=IFERROR(IF({Y}>0,INDEX(项_名称,{Y}),{tx("收费项")}),"")',
        '收付人认': f'=IFERROR(IF({Z}>0,INDEX(人_名称,{Z}),{tx("收付人")}),"")',
    }


# —— 第 1 行红字、第 3 行汇总条、_参 里录入本自己的几个键 ——
A_MOVED = '※ 表头或列挪了（台帐表头上面、基础资料里插删了行，或者插删了列）：查询本按位置读，会读错。马上撤销（Ctrl+Z）'
A_FULL = '※ 台帐满了：年底开新本；年中满了找我们放大'
A_NOFX = ('"※ 有 "&P_插行&" 行没有自动列（中间插了行，或整行按了 Delete）：这几行看不到余额，后面的余额也不对。'
          '漏记的录在最下面就行；改法：点行号选中上面一行复制，粘到这几行上，再改成要录的内容"')


def alert_formula():
    c = chain([
        ('P_挪了=1', f'"{A_MOVED}"'),
        ('P_插行>0', A_NOFX),
        ('P_可录<=0', f'"{A_FULL}"'),
        (f'P_可录<{NEAR}', '"△ 台帐还能录约 "&P_可录&" 行：年底开新本；年中满了找我们放大"'),
    ], '""')
    return f'=IFERROR({c},"{A_MOVED}")'


def entry_par_rows():
    """录入本 _参 自己的键（接在 PAR_KEYS 后面）：[(键, 公式, 说明)]"""
    L, B = q(SH_LEDGER), q(SH_BASE)
    hs = f'${ENTRY_HELP["账户号"]}'
    moved = (f'=IFERROR(IF(OR(ROW({L}!$A${HDR})<>{HDR},COLUMN({L}!$A${HDR})<>1,COLUMNS({L}!$A${HDR}:${LEDGER_LAST}${HDR})<>{CI(LEDGER_LAST)},'
             f'ROW({B}!$A${BASE_END_ROW})<>{BASE_END_ROW},COLUMN({B}!$A$6)<>1,COLUMNS({B}!$A$6:$N$6)<>14),1,0),1)')
    return [
        ('挪了', moved, '1＝资金台帐表头上面、基础资料里插删了行，或者插删了列（查询本按位置读会读错）'),
        ('插行', f'=IFERROR(ROWS({L}!{hs}${HDR}:{hs}${LEDGER_LAST_ROW})-COUNTA({L}!{hs}${HDR}:{hs}${LEDGER_LAST_ROW}),0)',
         '台帐里没有自动列的行数（中间插了行、或整行按了 Delete）'),
    ]


def kpi_formulas():
    return {
        '最后一笔': f'=IF(MAX({NM["日期值"]})>0,MAX({NM["日期值"]}),"")',
        '流水笔数': '=P_已用行数',
        '资金总额': f'=ROUND(P_期初合计+SUM({BLK_NET}),2)',
        '※ 行数': '=COUNTIF(台帐_问题,"※*")',
        '△ 行数': '=COUNTIF(台帐_问题,"△*")',
        '还能录': '=P_可录&" 行"',
    }


KPI_FMT = {'最后一笔': DATE, '流水笔数': '#,##0', '资金总额': MONEY, '※ 行数': '#,##0', '△ 行数': '#,##0', '还能录': 'General'}
KPI_NOTE = f'录的时候只开这本（打字不卡）。看报表、子表、核对表：先存盘，再打开同一个文件夹里的《{WB2_FILE[:-5]}》'


# —— _块：余额块表 ——
def build_blk(ws):
    a0 = CI(ENTRY_BLK['净'])
    p0 = CI(ENTRY_BLK['前'])
    ws[f'{ENTRY_BLK["有数"]}1'] = '块·有数'
    for a in range(1, N_ACC + 1):
        ws.cell(row=1, column=a0 + a - 1, value=f'块净·{a}')
        ws.cell(row=1, column=p0 + a - 1, value=f'块前·{a}')
    s = f'(ROW()-2)*{EB}+{R0 - HDR + 1}'
    e = f'(ROW()-2)*{EB}+{R0 - HDR + EB}'
    rng = lambda k: f'INDEX({NM[k]},{s}):INDEX({NM[k]},{e})'
    for b in range(1, ENTRY_N_BLK + 1):
        r = b + 1
        ws[f'{ENTRY_BLK["有数"]}{r}'] = f'=SUMPRODUCT(ABS({rng("净")}))'
        for a in range(1, N_ACC + 1):
            cn, cp = CL(a0 + a - 1), CL(p0 + a - 1)
            ws[f'{cn}{r}'] = f'=IF(${ENTRY_BLK["有数"]}{r}=0,0,SUMIFS({rng("净")},{rng("账户号")},COLUMN()-{a0 - 1}))'
            ws[f'{cp}{r}'] = '=0' if b == 1 else f'={cp}{r - 1}+{cn}{r - 1}'
    for row in ws.iter_rows(min_row=1, max_row=1):
        for c in row:
            c.font = F_HELP
    ws.sheet_state = 'hidden'


def blk_names():
    a0 = CI(ENTRY_BLK['净'])
    p0 = CI(ENTRY_BLK['前'])
    last = ENTRY_N_BLK + 1
    return {BLK_NET: f"{SH_BLK_E}!${CL(a0)}$2:${CL(a0 + N_ACC - 1)}${last}",
            BLK_PRE: f"{SH_BLK_E}!${CL(p0)}$2:${CL(p0 + N_ACC - 1)}${last}"}


# —— 首页（轻）——
HOME_ROWS = dict(alert=3, kpi_l=4, kpi_v=5, sec=7, hdr=8, first=9)


def build_home(ws, ctx):
    widths(ws, dict(A=7, B=30, C=14, D=16, E=16, F=16, G=3, H=22))
    ws.sheet_view.zoomScale = 90
    tip = ('录的时候只开这本：资金台帐、基础资料都在这本里，打字不卡。报表、子表、核对表在同一个文件夹的《' + WB2_FILE[:-5] +
           '》里，从这本自动取数：录完先存盘，再打开查询本（提示「更新 / 启用内容」时点它）。两本别同时开着录，文件名别改。')
    title(ws, '=IF(P_集团名称="","集团",P_集团名称)&" 资金台帐录入本"', 'F', tip, h2=46)
    R = HOME_ROWS
    ws.merge_cells(f'A{R["alert"]}:F{R["alert"]}')
    c = put(ws, f'A{R["alert"]}', f'={q(SH_LEDGER)}!$F$1', Font(name=YH, sz=11, bold=True, color=RGB_ERR_FONT), FILL_NONE,
            align=ALW, border=False)
    link(c, SH_LEDGER, 'A1')
    ws.conditional_formatting.add(f'A{R["alert"]}:F{R["alert"]}', FormulaRule(
        formula=[f'LEN($A${R["alert"]})>0'], font=Font(name=YH, sz=11, bold=True, color=RGB_ERR_FONT), fill=cf_bg(RGB_ERR_FILL)))
    ws.row_dimensions[R['alert']].height = 30
    kf = kpi_formulas()
    cells = list(zip('ABCDEF', ['资金总额', '流水笔数', '最后一笔', '※ 行数', '△ 行数', '还能录']))
    lbl = {'资金总额': '资金总额\n（期初＋全部流水）', '流水笔数': '流水笔数\n（台帐已用行数）', '最后一笔': '最后一笔\n（日期最晚的）',
           '※ 行数': '※ 行数\n（必须改）', '△ 行数': '△ 行数\n（看一眼）', '还能录': '台帐还能录'}
    for col, k in cells:
        put(ws, f'{col}{R["kpi_l"]}', lbl[k], F_KPI_L, FILL_KPI, align=ACW)
        put(ws, f'{col}{R["kpi_v"]}', kf[k], F_KPI_V, FILL_KPI, fmt=KPI_FMT[k], align=AC)
    link(ws[f'D{R["kpi_v"]}'], SH_LEDGER, f'{LEDGER_SHOW["问题"]}{HDR}')
    link(ws[f'E{R["kpi_v"]}'], SH_LEDGER, f'{LEDGER_SHOW["问题"]}{HDR}')
    ws.row_dimensions[R['kpi_l']].height = 32
    ws.row_dimensions[R['kpi_v']].height = 26
    ws.merge_cells(f'A{R["sec"]}:F{R["sec"]}')
    put(ws, f'A{R["sec"]}', '各账户当前余额（期初＋台帐里这个账户全部流水；跟网银余额对一下）', F_SEC, fill(C_HOME), align=AL)
    header(ws, R['hdr'], [('A', '序号'), ('B', '账户'), ('C', '公司'), ('D', '期初余额'), ('E', '本本收支净额'), ('F', '当前余额')], '首页',
           height=24)
    r0 = R['first']
    a = f'(ROW()-{r0 - 1})'
    ok = f'INDEX(账户_有效,{a})<>1'
    for i in range(N_ACC):
        r = r0 + i
        cono = f'INDEX(账户_公司号,{a})'
        put(ws, f'A{r}', f'=IF({ok},"",{a})', F_AUTO, FILL_NONE, align=AC)
        put(ws, f'B{r}', f'=IF({ok},"",INDEX(账户_名称,{a}))', F_TXT, FILL_NONE, align=AL)
        put(ws, f'C{r}', f'=IF({ok},"",IF({cono}=0,"",INDEX(公司_名称,{cono})))', F_TXT, FILL_NONE, align=AC)
        put(ws, f'D{r}', f'=IF({ok},"",INDEX(账户_期初,{a}))', F_TXT, FILL_NONE, fmt=MONEY, align=AR)
        put(ws, f'E{r}', f'=IF({ok},"",ROUND(SUM(INDEX({BLK_NET},0,{a})),2))', F_TXT, FILL_NONE, fmt=MONEY, align=AR)
        put(ws, f'F{r}', f'=IF({ok},"",ROUND(INDEX(账户_期初,{a})+SUM(INDEX({BLK_NET},0,{a})),2))', F_AUTOB, FILL_NONE, fmt=MONEY,
            align=AR)
        if i >= max(len(ctx.get('accounts') or []), 1) + 5:
            ws.row_dimensions[r].hidden = True
    rt = r0 + N_ACC
    put(ws, f'A{rt}', None, F_KPI_L, FILL_KPI)
    put(ws, f'B{rt}', '合计', F_KPI_L, FILL_KPI, align=AC)
    put(ws, f'C{rt}', None, F_KPI_L, FILL_KPI)
    for col in 'DEF':
        put(ws, f'{col}{rt}', f'=SUM({col}{r0}:{col}{rt - 1})', F_KPI_V, FILL_KPI, fmt=MONEY, align=AR)
    cf_neg_red(ws, f'F{r0}:F{rt}', f'F{r0}')
    put(ws, f'B{rt + 1}', '加了账户的话：选中上面藏起来的行，右键「取消隐藏」', F_NOTE, FILL_NONE, align=AL, border=False)
    # 导航
    put(ws, f'H{R["sec"]}', '跳到', F_SEC, fill(C_HOME), align=AC)
    for i, sh in enumerate([SH_LEDGER, SH_BASE, SH_HELP]):
        c = put(ws, f'H{R["hdr"] + i}', sh, F_LINK, FILL_NONE, align=AC)
        link(c, sh, 'A1')
    ws.freeze_panes = f'A{R["hdr"] + 1}'
    print_setup(ws, None, landscape=False)
    ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                    formatCells=False, selectLockedCells=False, selectUnlockedCells=False)


# —— 使用说明（两本各一份；大白话、短）——
def _help_entry(ctx):
    lc, ls = LEDGER_COLS, LEDGER_SHOW
    ins = '、'.join(['A～G', 'I～K', 'P'])
    return [
        ('一、这本是干什么的', [
            ('只在这里录', f'这本是录入本：每天的流水录在【资金台帐】，账户、公司、收费项、收付人、摘要规则在【基础资料】。'
                       f'报表、子表、核对表都在同一个文件夹的《{WB2_FILE[:-5]}》里，从这本自动取数，那本里不用录。'),
            ('首页', '资金总额、各账户当前余额、要改的 ※ 和最好补的 △ 有几行。'),
        ]),
        ('二、每天怎么录', [
            ('一笔一行', '从最下面第一个空行接着录。日期、账户必填（淡黄色）；收入、支出只填一个。'),
            ('收费项', f'最好从下拉选；空着会按摘要认（{ls["收费项认"]} 列琥珀色斜体）。一笔钱拆几行的一定要选。'),
            ('内部转账', f'集团里两个账户互转记两行（转出那户记支出、转入那户记收入），「转账的另一户」（{lc["对方账户"]} 列）都填对方。'),
            ('灰色几列', f'{ls["余额"]} 余额、{ls["问题"]} 问题、{ls["属性"]} 属性、{ls["收费项认"]}/{ls["收付人认"]} 认出来的收费项和收付人，自动算，别动。'
                      f'{ls["问题"]} 列 ※ 必须改，△ 看一眼。'),
            ('这本查什么', '这本只查「这一行」看得出来的毛病（日期、账户、金额、名字认不认得……）。内部转账配没配上、日期跳期、重复粘贴、'
                      '日期顺序这些要看整张表的，在查询本的核对表里看。'),
            ('粘网银', f'用「选择性粘贴 → 数值」，只粘 {ins} 这几列，别粘到灰色的列上。'),
        ]),
        ('三、别插行', [
            ('漏记的', '直接记在最下面，日期写对就行：查询本的报表、子表都按日期算，跟上下顺序没关系。'),
            ('调顺序', '用「日期」列表头的筛选箭头排序（升序）。别选中整张表排序。'),
            ('不要的行', f'点左边行号选整行，右键「删除」。只想清内容：只选 {ins} 这几列按 Delete，别整行按 Delete（会把右边的自动列也清掉）。'),
            ('插了怎么办', '第 1 行会红字「有 x 行没有自动列」：马上撤销（Ctrl+Z）；已经录了，点行号选中上面一行复制，粘到插进来的行上，再改成要录的内容。'),
            ('别动的', '别在表头上面插删行、别插删列（查询本按位置读，会读错）；别用「删除重复项」，别转成「表格」。'),
        ]),
        ('四、基础资料', [
            ('往下加', '账户、公司、收费项、收付人、规则都在各块下面的空行接着加；别在中间插行、删行（查询本按位置读）。③账户可以排序。'),
            ('改名', '改名时把旧名字加进「别名」（多个用「、」隔开），台帐里以前的写法照样认。不用的账户把状态改「销户」，填停用日，别删行。'),
            ('录入模式', '两本版不用「录入模式」：录的时候只开这本，本来就不卡。'),
        ]),
        ('五、存盘、两本一起用', [
            ('存盘', '录完按 Ctrl+S 存盘。查询本取的是这本存过盘的数。'),
            ('同一个文件夹', f'《{WB1_FILE[:-5]}》《{WB2_FILE[:-5]}》放在同一个文件夹，文件名别改（查询本按文件名找这本）。'),
            ('别同时开着录', '录的时候别开着查询本：两本都开着时，这边每录一格那边也跟着算，又会卡。'),
        ]),
    ]


def _help_query(ctx, one_book_sections):
    keep = [s for s in one_book_sections if s[0].startswith(('六、', '七、', '十三、'))]
    ren = {'六、': '五、', '七、': '六、', '十三、': '七、'}
    keep = [(next((v + t[len(k):] for k, v in ren.items() if t.startswith(k)), t), items) for t, items in keep]
    return [
        ('一、这本是干什么的', [
            ('只看不录', f'这本是查询本：首页、报表、子表、核对表。数全从同一个文件夹的《{WB1_FILE[:-5]}》来，这本里不录流水、不改基础资料'
                       '（隐藏的「资金台帐」「基础资料」是取数用的，别动）。能改的只有报表顶上的黄格（年份、日期、月份、账户）。'),
        ]),
        ('二、打开以前', [
            ('先存录入本', '在录入本里录完、存盘，再打开这本。'),
            ('点更新', 'Excel 打开时顶上提示「安全警告：已禁用外部数据连接」或弹窗问要不要更新，点「启用内容」或「更新」；'
                     'WPS 点「启用内容 / 更新链接」。没提示的话：数据 → 编辑链接 → 更新值。不点就是上次存盘时的数。'),
            ('对一下', '首页第 3 行右边「数据来自录入本：流水 x 笔，最后一笔 …」跟录入本首页的「流水笔数」「最后一笔」一样，就是连上最新的了。'),
        ]),
        ('三、两本放一起', [
            ('同一个文件夹', f'《{WB1_FILE[:-5]}》《{WB2_FILE[:-5]}》放在同一个文件夹，文件名别改。'),
            ('改了名', '改了录入本的名字、挪了文件夹、从微信另存多了「(1)」：在这本点「数据 → 编辑链接 → 更改源」，选录入本。'),
            ('别同时开着录', '录的时候只开录入本。两本都开着时，录入本每录一格这本也跟着算，又会卡。'),
        ]),
        ('四、首页红字', [
            ('「录入本的表头挪了」', '录入本资金台帐表头上面、或基础资料里插删了行、插删了列：这本按位置读，读错了。回录入本撤销，'
                             '基础资料别在中间插行（在各块下面空行加）。'),
            ('别的红字、核对表', '看下面「七、核对表怎么看」；台帐第几行说的是录入本资金台帐的行号。'),
        ]),
    ] + keep


def build_help_sheet(ws, ttl, tip, sections):
    """照一本版 s_check.build_help 的版式写一张使用说明"""
    from s_check import _lines, WRAP_TOP
    widths(ws, dict(A=4, B=16, C=110))
    title(ws, ttl, 'C', tip, h2=24)
    home_link(ws, 'A3')
    put(ws, 'B3', '目录', F_KPI_L, FILL_KPI, align=AC)
    put(ws, 'C3', '　'.join(t for t, _i in sections), F_NOTE, FILL_NONE, align=ALW)
    ws.row_dimensions[3].height = 30
    r = 5
    hf = fill('FF1F3864')
    for t, items in sections:
        ws.merge_cells(f'A{r}:C{r}')
        put(ws, f'A{r}', t, F_SEC, hf, align=AL)
        ws.row_dimensions[r].height = 22
        r += 1
        for i, (k, v) in enumerate(items, 1):
            put(ws, f'A{r}', i, F_AUTO, FILL_NONE, align=Alignment(horizontal='center', vertical='top'))
            put(ws, f'B{r}', k, F_TXTB, FILL_NONE, align=WRAP_TOP)
            put(ws, f'C{r}', v, F_TXT, FILL_NONE, align=WRAP_TOP)
            ws.row_dimensions[r].height = max(18, 14 * max(_lines(v, 110), _lines(k, 16)) + 6)
            r += 1
        r += 1
    ws.merge_cells(f'A{r}:C{r}')
    put(ws, f'A{r}', '有不明白的、数对不上的，把首页、核对表截图发给我们。', F_NOTE, FILL_NONE, align=AL, border=False)
    ws.freeze_panes = 'A5'
    print_setup(ws, None, landscape=False)
    ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                    formatCells=False, selectLockedCells=False, selectUnlockedCells=False)


def build_help_entry(ws, ctx):
    build_help_sheet(ws, '使用说明（录入本版）', '一共五条，先看「三、别插行」和「五、存盘、两本一起用」。', _help_entry(ctx))


def build_help_query(ws, ctx):
    from s_check import _help_text
    build_help_sheet(ws, '使用说明（查询本版）', '先看「二、打开以前」；报表怎么看见「五」～「七」。', _help_query(ctx, _help_text(ctx)))
