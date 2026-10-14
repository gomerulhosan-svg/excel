# -*- coding: utf-8 -*-
"""从 2026年日记账 读出 24 张账户表（conf.ACCOUNTS）的流水（缓存值），不做任何清洗。"""
import openpyxl
from openpyxl.utils import column_index_from_string as CI

from conf import LAYOUT, ACCOUNTS


def _num(v):
    if v is None or v == '':
        return 0.0
    if isinstance(v, bool):
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace(',', '').strip()
    try:
        return float(s)
    except ValueError:
        return None


def _txt(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v)


def read_journal(src, sheets_meta):
    """返回 (rows, openings, skipped)。
    rows: 每个流水行一条 dict；openings: {原表: (期初数, 依据)}；skipped: 有内容但没金额的行。"""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    rows, openings, skipped = [], {}, []
    for sheet, *_ in ACCOUNTS:
        lay = LAYOUT[sheet]
        meta = sheets_meta[sheet]
        start, end = meta['数据起始行'], meta['最后流水行']
        sumrows = set(meta['行数']['本月合计行号'])
        ws = wb[sheet]
        data = list(ws.iter_rows(min_row=1, max_row=end, max_col=26, values_only=True))

        def g(r, col):
            return data[r - 1][CI(col) - 1] if col else None

        # 期初
        orow = meta['期初行']
        if orow:
            openings[sheet] = (round(_num(g(orow, lay['bal'])), 2), '%s!%s%d' % (sheet.strip(), lay['bal'], orow))
        else:
            openings[sheet] = (0.0, '原表没有期初行')
        for r in range(start, end + 1):
            if r in sumrows or r == orow:
                continue
            row = data[r - 1]
            texts = [_txt(x).strip() for x in row[1:7] if isinstance(x, str)]
            if any(t in ('本月合计', '期初余额') for t in texts):
                continue
            m, d = g(r, 'B'), g(r, 'C')
            rec = dict(sheet=sheet, row=r,
                       m=int(m) if isinstance(m, (int, float)) else None,
                       d=int(d) if isinstance(d, (int, float)) else None,
                       item=_txt(g(r, lay.get('item'))).strip(),
                       party_raw=_txt(g(r, lay['party'])),
                       memo_raw=_txt(g(r, lay['memo'])),
                       note=_txt(g(r, lay.get('note'))).strip(),
                       owner=_txt(g(r, lay.get('owner'))).strip(),
                       bal=_num(g(r, lay['bal'])),
                       memocol=lay['memo'])
            rec['extra'] = {nm: _txt(g(r, c)).strip() for c, nm in lay.get('extra', {}).items()
                            if _txt(g(r, c)).strip()}
            bad = []
            for k in ('inc', 'exp', 'fee', 'gross'):
                col = lay.get(k)
                v = _num(g(r, col)) if col else 0.0
                if v is None:
                    bad.append('%s%d=%r' % (col, r, g(r, col)))
                    v = 0.0
                rec[k] = round(v, 2)
            cats = {}
            for col, nm in lay.get('cats', {}).items():
                v = _num(g(r, col))
                if v is None:
                    bad.append('%s%d=%r' % (col, r, g(r, col)))
                    continue
                if abs(v) > 1e-9:
                    cats[nm] = round(v, 2)
            rec['cats'] = cats
            rec['catcols'] = {nm: col for col, nm in lay.get('cats', {}).items()}
            rec['cols'] = {k: lay.get(k) for k in ('item', 'party', 'memo', 'inc', 'exp', 'fee', 'gross')}
            if bad:
                rec['bad'] = bad
            has_amt = rec['inc'] or rec['exp'] or cats or rec['fee'] or rec['gross']
            if not has_amt:
                if rec['memo_raw'].strip() or rec['party_raw'].strip() or rec['item']:
                    skipped.append(rec)
                continue
            rows.append(rec)
    wb.close()
    return rows, openings, skipped


def read_guwan_bank_opening(src):
    """古玩城农商银行 银行口径期初：R5（用户手算列）。"""
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb['古玩城农商银行']
    v = None
    for row in ws.iter_rows(min_row=5, max_row=5, min_col=18, max_col=18, values_only=True):
        v = row[0]
    wb.close()
    return round(float(v), 2)



def _used_cols(lay):
    used = {'A', 'B', 'C'}
    for k, v in lay.items():
        if isinstance(v, str) and 1 <= len(v) <= 2 and v.isalpha() and v.isupper():
            used.add(v)
    used |= set(lay.get('cats', {})) | set(lay.get('extra', {}))
    return used


def read_stray(src, sheets_meta):
    """不导的零散数：返回 [(原表, 单元格, 数值, 种类)]。
    种类 残留＝读懂报告标的「残留行」（数据区下面）；表外＝流水区里、不在本表版式用到的列里的数（草稿、手算列）。"""
    from openpyxl.utils import get_column_letter
    out = []
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    for sheet, *_ in ACCOUNTS:
        meta = sheets_meta[sheet]
        used = _used_cols(LAYOUT[sheet])
        res = set(meta['行数'].get('残留行') or [])
        start, end = meta['数据起始行'], meta['最后流水行']
        ws = wb[sheet]
        for r, row in enumerate(ws.iter_rows(min_row=1, max_row=max([end] + list(res)), max_col=26, values_only=True), 1):
            if r in res:
                for c, v in enumerate(row, 1):
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        out.append((sheet, '%s%d' % (get_column_letter(c), r), v, '残留'))
            elif start <= r <= end:
                for c, v in enumerate(row, 1):
                    col = get_column_letter(c)
                    if col in used or v is None or v == '':
                        continue
                    out.append((sheet, '%s%d' % (col, r), v, '表外'))
    wb.close()
    return out
