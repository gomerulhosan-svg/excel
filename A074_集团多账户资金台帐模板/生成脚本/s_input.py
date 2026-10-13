# -*- coding: utf-8 -*-
"""录入（s_input）：【资金台帐】（录入列＋有效性＋输入提示、条件格式、显示列 H/L/M/N/O、自动筛选、第 3 行汇总条、导入的流水）
   和【基础资料】（① 参数 ② 公司 ③ 账户 ④ 收费项 ⑤ 收付人 ⑥ 摘要关键词规则：序号、核对列、有效性、预置数据）。设计见 设计.md §3、§4。

   资金台帐只放录入列＋5 个显示列（余额 H、问题 L、属性 M、收费项认 N、收付人认 O）。显示列只是 INDEX(流_xx, 第几条) 给人看，
   别的公式不引用它们（核对表「显示列被盖掉」除外）。第 n 条＝ROW()-ROW($A$4)，插行、删行、排序都对得上。
   12,000 行的录入格用「列样式」（<col style>）上色，只有写了值的格子逐格设样式（文件小）。"""
import datetime as dt
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.dimensions import SheetFormatProperties
from layout import *
from common import *

# ───────────────────────── 小工具 ─────────────────────────
F_DARK = Font(name=YH, sz=10, color='FF404040')     # 生成时写死、别改的格（子表名）


def _val(v):
    """写进格子的值：date → datetime（openpyxl 按日期存）；空串、None → None"""
    if v is None or (isinstance(v, str) and v == ''):
        return None
    if isinstance(v, dt.datetime):
        return v
    if isinstance(v, dt.date):
        return dt.datetime(v.year, v.month, v.day)
    return v


def _set(ws, r, c, v, style=None):
    """写一格（字符串以 = 开头的当文字，不当公式）"""
    v = _val(v)
    if v is None:
        cell = ws.cell(row=r, column=c)
    else:
        cell = ws.cell(row=r, column=c, value=v)
        if isinstance(v, str) and v.startswith('='):
            cell.data_type = 's'
    if style:
        _apply(cell, style)
    return cell


def _apply(cell, st):
    font, fill_, fmt, align = st
    cell.font, cell.fill, cell.border, cell.alignment = font, fill_, BD, align
    if fmt:
        cell.number_format = fmt


def _col_style(ws, col, st):
    """整列的列样式：没写值的空格子也是这个样子（Excel / WPS / LibreOffice 都认 <col style>）"""
    font, fill_, fmt, align = st
    cd = ws.column_dimensions[col]
    cd.font, cd.fill, cd.border, cd.alignment = font, fill_, BD, align
    if fmt:
        cd.number_format = fmt


def dv_add(ws, sqref, *, type=None, formula1=None, formula2=None, operator=None, stop=True, prompt=None, title='提示',
           err_title=None, err=None, show_err=True):
    dv = DataValidation(type=type, formula1=formula1, formula2=formula2, operator=operator, allow_blank=True,
                        showErrorMessage=show_err and type is not None, errorStyle='stop' if stop else 'warning')
    if prompt:
        assert len(prompt) <= 255, prompt
        dv.promptTitle, dv.prompt, dv.showInputMessage = title, prompt, True
    if err:
        assert len(err) <= 255, err
        dv.errorTitle, dv.error = err_title or '填得不对', err
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def dv_fixed(ws, sqref, key, prompt=None, stop=True, err=None):
    return dv_list(ws, sqref, '"' + DV_FIXED[key] + '"', prompt=prompt, stop=stop,
                   error=err or f'只能从下拉里选：{DV_FIXED[key].replace(",", "、")}')


def _in_list(x, items):
    """x 不空而且不在清单里（清单直接写成数组常量）"""
    arr = '{' + ','.join(f'"{t}"' for t in items) + '}'
    return f'AND({raw(x)},ISERROR(MATCH(TRIM({x}&""),{arr},0)))'


def _cnt_dup(pre, x):
    """名字（规范写法）在 _表 名称规 里出现不止一次 → 重名"""
    return f'COUNTIF({pre}_名称规,"="&{esc(norm(x))})>1'


def _any_raw(cells):
    return 'OR(' + ','.join(raw(c) for c in cells) + ')'


# ═════════════════════════════ 资金台帐 ═════════════════════════════
LED_TIP = (
    '★ 所有账户的流水都录在这一张表，一笔一行，日期、账户必填（淡黄）。收入、支出只填一个；网银导出的「交易金额 −5000」照样粘进收入栏，会自动当支出。'
    '内部转账（集团里两个账户互转）记两行：转出那户记支出、转入那户记收入，两行的「转账的另一户」都填对方账户，收费项空着会自动当内部转账。'
    '收费项空着会按摘要里的关键词去认（N 列琥珀色斜体就是认出来的）；一笔钱拆成几行的一定要自己选。'
    '例：C118 交 1 月租金 1,500＋物业 150＋水电 105.6＝1,755.6，录 3 行（选中上一行 A:E 按 Ctrl+D 复制，再改收费项和金额），'
    '银行回单上的 1,755.6 对的是这 3 行合计；手续费另起一行选「银行手续费」。'
    '粘网银流水用「选择性粘贴→数值」，只粘到 A～G、I～K 列，文字日期、带逗号的金额都认。'
    '灰色几列（H、L～O）是自动算的，别动。L 列 ※ 必须改，△ 看一眼。排序点日期列的筛选箭头。'
    '月底核对完可以「固化」：先把 ※ 改完，再把 N 列复制、选择性粘贴数值到 C 列（O 列同样粘到 D 列）。'
    '一次录很多时，去【基础资料】①把「报表暂停」填「是」，录完清空。')

# 日期列：真日期照 yyyy/mm/dd 显示；网银粘进来的 8 位数（20261001）照原样显示，不显示成 ######
DATE_IN = '[<73051]yyyy/mm/dd;0'
# 每列：(字体, 底色, 数字格式, 对齐)
_LEDGER_STYLE = {
    '日期': (F_IN, FILL_IN, DATE_IN, AC), '账户': (F_IN, FILL_IN, None, AL),
    '收费项': (F_IN, FILL_NONE, None, AL), '收付人': (F_IN, FILL_NONE, None, AL), '摘要': (F_IN, FILL_NONE, None, AL),
    '收入': (F_IN, FILL_NONE, MONEY, AR), '支出': (F_IN, FILL_NONE, MONEY, AR),
    '对方账户': (F_IN, FILL_NONE, None, AL), '归属公司': (F_IN, FILL_NONE, None, AC), '备注': (F_IN, FILL_NONE, None, AL),
    '原表位置': (F_IN, FILL_NONE, None, AL),
    '余额': (F_AUTO, FILL_AUTO, MONEY, AR), '问题': (F_AUTO, FILL_AUTO, None, AL), '属性': (F_AUTO, FILL_AUTO, None, AC),
    '收费项认': (F_AUTO, FILL_AUTO, None, AL), '收付人认': (F_AUTO, FILL_AUTO, None, AL),
}
_N = f'ROW()-ROW($A${HDR})'                         # 这一行是第几条
_ATTR_NAMES = {v: k for k, v in ATTR_NO.items()}


def ledger_show_formulas():
    """显示列公式（每行一样：相对位置只有 ROW()，排序、插行都不坏）"""
    n = _N
    choose = ','.join(f'"{_ATTR_NAMES.get(p, "")}"' for p in range(1, max(ATTR_NO.values()) + 1))
    return {
        '余额': f'=IFERROR(INDEX(流_余额,{n}),"")',
        '问题': f'=IFERROR(INDEX(流_问题,{n})&"","")',
        '属性': f'=IFERROR(IF(INDEX(流_有效,{n})<>1,"",CHOOSE(INDEX(流_属性号,{n}),{choose})),"")',
        '收费项认': f'=IFERROR(IF(INDEX(流_项号,{n})>0,INDEX(项_名称,INDEX(流_项号,{n})),""),"")',
        '收付人认': f'=IFERROR(IF(INDEX(流_人号,{n})>0,INDEX(人_名称,INDEX(流_人号,{n})),""),"")',
    }


def ledger_kpi_formulas():
    """第 3 行汇总条（本年＝管理年度；经营收入＝属性「收入」的行 Σ 净，经营支出＝属性「支出」的行 −Σ 净）"""
    y = 'P_管理年度'
    yr = f'流_月码,">="&({y}*100+1),流_月码,"<="&({y}*100+12)'
    return {
        '管理年度': f'={y}',
        '本年经营收入': f'=ROUND(SUMIFS(流_净,{yr},流_属性号,{A_IN}),2)',
        '本年经营支出': f'=ROUND(-SUMIFS(流_净,{yr},流_属性号,{A_OUT}),2)',
        '本年经营净额': None,           # ＝收入格－支出格（下面按格子拼）
        '内部转账净额': f'=ROUND(SUMIFS(流_净,{yr},流_属性号,{A_INT}),2)',
        '资金总额': '=ROUND(P_期初合计+SUM(流_净),2)',
        '已用行数': f'=P_已用行数&" / {N_LEDGER}"&IF(P_已用行数>={int(N_LEDGER * 0.9)},"  △ 快满了：开新本","")',
        '※ / △': '=IFERROR("※ "&N(核对_问题)&" 行　△ "&N(核对_提醒)&" 行","")',
    }


def build_ledger(ws, ctx):
    C, S = LEDGER_COLS, LEDGER_SHOW
    allc = {**C, **S}
    last = LEDGER_LAST_ROW
    widths(ws, LEDGER_WIDTH)
    title(ws, '=IF(P_集团名称="","集团",P_集团名称)&" 资金台帐（"&P_开账年&" 年，开账 "&YEAR(P_开账日)&"/"&MONTH(P_开账日)'
              '&"/"&DAY(P_开账日)&"）"', LEDGER_LAST, LED_TIP)
    # 第 3 行汇总条
    kf = ledger_kpi_formulas()
    lbl = {'内部转账净额': '本年内部转账净额\n（应该是 0）', '已用行数': '已用行数 / 容量', '※ / △': '※ / △\n（看 L 列）'}
    cells = {k: (a, b) for a, b, k in LEDGER_KPI}
    for a, b, k in LEDGER_KPI:
        f = kf[k]
        if k == '本年经营净额':
            f = f'={cells["本年经营收入"][1]}-{cells["本年经营支出"][1]}'
        fmt = FMT_YEAR if k == '管理年度' else (MONEY if k not in ('已用行数', '※ / △') else '@')
        kpi(ws, a, b, lbl.get(k, k), f, fmt=fmt)
        if k in ('已用行数', '※ / △'):
            ws[b].font = Font(name=YH, sz=10, bold=True, color='FFC00000')
            ws[b].alignment = ACW
    ws.row_dimensions[3].height = 35
    cf_neg_red(ws, f'{cells["本年经营净额"][1]}', cells['本年经营净额'][1])
    # 第 4 行表头（录入列、自动列同一个底色；自动列表头写「自动」）
    header(ws, HDR, [(allc[k], LEDGER_HEAD[k]) for k in allc], '录入', height=36)
    # 列样式（12,000 行的空格子不逐格写）
    for k, col in allc.items():
        _col_style(ws, col, _LEDGER_STYLE[k])
    ws.sheet_format = SheetFormatProperties(defaultRowHeight=16.5, customHeight=True)
    # 导入的流水：照 entries 顺序写进录入列（日期写真日期、金额写数，0 写空）
    entries = ctx.get('entries') or []
    keys = list(C)
    for i, e in enumerate(entries[:N_LEDGER]):
        r = R0 + i
        e = e or {}
        for k in keys:
            v = e.get(k)
            if k in ('收入', '支出') and isinstance(v, (int, float)) and v == 0:
                v = None
            _set(ws, r, CI(C[k]), v, _LEDGER_STYLE[k])
    # 显示列（每行一样的公式，share.py 会压成共享公式）
    sf = ledger_show_formulas()
    for r in range(R0, last + 1):
        for k, col in S.items():
            c = ws.cell(row=r, column=CI(col), value=sf[k])
            _apply(c, _LEDGER_STYLE[k])
    # 有效性、输入提示
    rng = lambda k: f'{allc[k]}{R0}:{allc[k]}{last}'
    dv_add(ws, rng('日期'), type='date', operator='between', formula1='36526', formula2='73050', stop=True,
           prompt='写真日期：2026/10/1 或 2026-10-1。网银粘进来的文字日期也认（L 列提示 △）。年份、月份写错（比前面最晚的一笔晚 60 天以上、下一笔又早回去）会报 ※',
           err_title='日期', err='要填真正的日期，像 2026/10/1 或 2026-10-1。手打的「2026.10.1」「10.1」是文字或小数，填不进去')
    dv_list(ws, rng('账户'), '账户列表', stop=False,
            prompt='从下拉选账户；写原表名（如「古玩城农商银行」）也认。新开的卡先去【基础资料】③加一行；改了名把旧名加进③的别名')
    dv_list(ws, rng('收费项'), '收费项列表', stop=False,
            prompt='从下拉选（常用的在上面）。空着会按摘要关键词认（N 列琥珀色斜体）；一笔拆成几行的一定要选。手续费另起一行选「银行手续费」')
    dv_list(ws, rng('收付人'), '收付人列表', stop=False,
            prompt='从下拉选；空着会按摘要认。新的收付人先去【基础资料】⑤加；改了名把旧写法加进别名')
    dv_add(ws, rng('摘要'), prompt='照银行回单写。收费项、收付人空着时，按这里的关键词去认（规则在【基础资料】⑥）')
    for k in ('收入', '支出'):
        dv_add(ws, rng(k), type='decimal', operator='between', formula1='-999999999999', formula2='999999999999', stop=False,
               prompt='收入、支出只填一个；手续费另起一行选「银行手续费」。网银的负数照粘，会自动当另一栏',
               err_title='金额', err='金额要填数字。「10000+200」这种要写成 =10000+200；点「是」照样填进去（L 列会报 ※）')
    dv_list(ws, rng('对方账户'), '账户列表', stop=False,
            prompt='只有集团里两个账户互转才填（填对方那一户），收费项空着会自动当内部转账。外单位、个人请填在收付人',
            error='这里只填本册的账户（【基础资料】③）。外单位、个人请填在收付人；点「是」照样填进去（L 列会报 ※）')
    dv_list(ws, rng('归属公司'), '公司列表', stop=True,
            prompt='只有私户替别的公司收付才填：收支算到这家公司，余额还是算账户所属的公司',
            error='从下拉选【基础资料】②里的公司简称')
    dv_add(ws, rng('备注'), prompt='开票状态、项目、拆账说明都写这里（以后用筛选「包含」查）。收入类记在支出栏的退款，备注写「冲减」就不提示')
    dv_add(ws, rng('原表位置'), prompt='导入时记的原日记账位置（表名!格子）；新录的不用填')
    for k in S:
        dv_add(ws, rng(k), prompt='灰色是自动算的，别在这里打字、粘贴。中间插了行就选中上一行的 H、L～O 往下拖一格')
    # 条件格式（同表引用；先加的优先）
    cf_neg_red(ws, rng('余额'), f'{S["余额"]}{R0}')
    cf_flags(ws, rng('问题'), f'{S["问题"]}{R0}')
    cf_row_flags(ws, f'A{R0}:{LEDGER_LAST}{last}', f'${S["问题"]}{R0}')
    cf_guess(ws, rng('收费项认'), f'AND(TRIM(${C["收费项"]}{R0}&"")="",${S["收费项认"]}{R0}<>"")')
    cf_guess(ws, rng('收付人认'), f'AND(TRIM(${C["收付人"]}{R0}&"")="",${S["收付人认"]}{R0}<>"")')
    ws.freeze_panes = LEDGER_FREEZE
    ws.auto_filter.ref = LEDGER_FILTER
    print_setup(ws, f'{HDR}:{HDR}', landscape=True)


# ═════════════════════════════ 基础资料 ═════════════════════════════
BASE_TIP = (
    '★ 全书的名单和参数都在这里：① 参数 ② 公司 ③ 账户 ④ 收费项 ⑤ 收付人 ⑥ 摘要关键词规则（点第 3 行跳过去）。'
    '淡黄格可以改，灰色是自动的（序号、核对）。名单往各块下面的空行接着加；③可以插行、排序（子表靠「子表名」认账户，那一列别改）。'
    '账户、收费项、收付人改名时，把旧名字加进「别名」（用「、」隔开），台帐里旧的写法照样认。'
    '不用的账户把状态改成「销户」、填停用日，别删行。年中新开的账户期初余额填 0，开户那天的余额在台帐录一笔「账户转入」。'
    '核对列 ※ 必须改，△ 看一眼。')

def _block_rows(k):
    h = BASE_HDR[k]
    return h + 1, h + BASE_CAP[k]


def _check_formula(k, r):
    """各块核对列（※ 必须改、△ 看一眼；只报第一个）"""
    cols = BASE_COLS[k]
    c = {f: f'${col}{r}' for f, col in cols.items()}
    h = BASE_HDR[k]
    kpos = f'(ROW()-ROW($A${h}))'
    key = c[{'公司': '公司简称', '账户': '账户名称', '收费项': '收费项', '收付人': '收付人', '规则': '关键词'}[k]]
    others = [c[f] for f in cols if f not in ('序号', '核对', '规则号', '说明') and c[f] != key]
    pairs = []
    if k == '公司':
        pairs = [(f'NOT({raw(key)})', f'IF({_any_raw(others)},"※ 公司简称没填：这一行不算","")'),
                 (_cnt_dup('公司', key), '"※ 简称重复了（空格、全半角括号不同也算重名）：改一个"')]
    elif k == '账户':
        bad_ch = 'OR(' + ','.join(f'ISNUMBER(FIND("{ch}",{key}))' for ch in '[]:*?/\\') + ')'
        op = numv(c['期初余额'])
        st = f'TRIM({c["状态"]}&"")'
        pairs = [
            (f'NOT({raw(key)})', f'IF({_any_raw(others)},"※ 账户名称没填：这一行不算","")'),
            (_cnt_dup('账户', key), '"※ 账户名称重复了：改一个（旧名放别名）"'),
            (f'NOT({raw(c["公司"])})', '"※ 公司没填：填②里的公司简称"'),
            (f'{name2no(c["公司"], "公司")}=0', '"※ 公司不在②：先去②加，或改成②里的简称"'),
            (f'NOT({raw(c["期初余额"])})', '"※ 期初余额没填：没有就填 0"'),
            (f'ISTEXT({c["期初余额"]})', '"※ 期初余额是文字：改成数字"'),
            (_in_list(c['户别'], ACC_KINDS), '"※ 户别不在清单：从下拉选"'),
            (_in_list(c['状态'], ACC_STATUS), '"※ 状态只能填 在用/暂停/销户（空＝在用）"'),
            (f'AND({raw(c["子表名"])},COUNTIF(账户_子表名规,"="&{esc(norm(c["子表名"]))})>1)',
             '"※ 子表名重复了：子表靠它认账户，这一列别改"'),
            (f'OR(LEN({norm(key)})>20,{bad_ch})', '"△ 名字超过 20 字或带 []:*?/\\：子表、报表显示不全"'),
            # 年中才开账（开账日晚于开账年的 1 月 1 日）期初还填了数：会把这个数算进开账前各月的资金。
            # 1 月 1 日开账的户期初≠0 是正常的（不常动的户开账后几个月没流水也不报）。不引用 _流。
            (f'AND({op}<>0,P_开账日>DATE(YEAR(P_开账日),1,1))',
             '"△ 开账日不是 1 月 1 日（年中才开账），期初应填 0：在台帐开账那天录一笔收费项「账户转入」，金额＝这个期初，再把期初改成 0"'),
            (f'AND(OR({st}="暂停",{st}="销户"),NOT({raw(c["停用日"])}))', '"△ 暂停/销户请填停用日（从哪天起不用）"'),
        ]
    elif k == '收费项':
        pairs = [
            (f'NOT({raw(key)})', f'IF({_any_raw(others)},"※ 收费项没填：这一行不算","")'),
            (_cnt_dup('项', key), '"※ 收费项重复了：改一个（旧写法放别名）"'),
            (_in_list(c['属性'], ITEM_ATTRS), '"※ 属性只能选：' + '/'.join(ITEM_ATTRS) + '"'),
            (f'AND({raw(c["大类"])},IFERROR(INDEX(项_大类号,{kpos}),0)=99)', f'"※ 大类超过 {N_GROUP} 种：合并几个，报表才单列"'),
            (f'NOT({raw(c["属性"])})', '"△ 属性没填：按方向算（收入栏算收入、支出栏算支出）"'),
            (f'NOT({raw(c["大类"])})', '"△ 大类没填：报表里不单列大类"'),
        ]
    elif k == '收付人':
        pairs = [
            (f'NOT({raw(key)})', f'IF({_any_raw(others)},"※ 收付人没填：这一行不算","")'),
            (_cnt_dup('人', key), '"※ 收付人重复了：改一个（旧写法放别名）"'),
            (_in_list(c['类型'], PARTY_TYPES), '"△ 类型不在清单：收支报表按类型合计时算不到"'),
        ]
    elif k == '规则':
        pairs = [
            (f'NOT({raw(key)})', f'IF({_any_raw(others)},"※ 关键词没填：这条规则不起作用","")'),
            (f'AND(NOT({raw(c["收费项"])}),NOT({raw(c["收付人"])}))', '"※ 收费项、收付人都没填：这条规则没用"'),
            (_in_list(c['查找位置'], RULE_WHERE), '"※ 查找位置只能填 摘要/收付人（空＝都找）"'),
            (_in_list(c['方向'], RULE_DIRS), '"※ 方向只能填 收/支/都（空＝都）"'),
            (f'AND({raw(c["适用范围"])},{name2no(c["适用范围"], "公司")}=0,{name2no(c["适用范围"], "账户")}=0)',
             '"※ 适用范围要填②的公司简称或③的账户名（空＝全部）"'),
            (f'AND({raw(c["收费项"])},{name2no(c["收费项"], "项")}=0)', '"※ 收费项不在④：从下拉选，或先去④加"'),
            (f'AND({raw(c["收付人"])},{name2no(c["收付人"], "人")}=0)', '"※ 收付人不在⑤：从下拉选，或先去⑤加"'),
        ]
    return '=' + chain(pairs, '""')


# 各块要填的列：(字体, 底色, 数字格式, 对齐)
_IN_TXT = (F_IN, FILL_IN, None, AL)
_IN_C = (F_IN, FILL_IN, None, AC)
_ST_SEQ = (F_AUTO, FILL_AUTO, None, AC)
_ST_CHK = (F_AUTO, FILL_AUTO, None, AL)
_BASE_STYLE = {
    '公司': dict(公司简称=_IN_TXT, 公司全称=_IN_TXT, 备注=_IN_TXT),
    '账户': dict(账户名称=_IN_TXT, 子表名=(F_DARK, FILL_AUTO, None, AL), 公司=_IN_C, 户别=_IN_C, 开户行=_IN_C, 户号尾号=_IN_C,
               持卡人=_IN_C, 期初余额=(F_IN, FILL_IN, MONEY, AR), 状态=_IN_C, 停用日=(F_IN, FILL_IN, DATE, AC), 别名=_IN_TXT,
               备注=_IN_TXT),
    '收费项': dict(收费项=_IN_TXT, 属性=_IN_C, 大类=_IN_C, 别名=_IN_TXT, 说明=_IN_TXT),
    '收付人': dict(收付人=_IN_TXT, 类型=_IN_C, 别名=_IN_TXT, 备注=_IN_TXT),
    '规则': dict(关键词=_IN_TXT, 查找位置=_IN_C, 方向=_IN_C, 适用范围=_IN_C, 收费项=_IN_C, 收付人=_IN_C, 规则号=_IN_C,
               说明=_IN_TXT),
}


def _preset_rows(k, ctx):
    """预置数据 → 每行 {字段: 值}（字段名＝BASE_COLS[k] 的键）"""
    if k == '公司':
        return [dict(公司简称=x['公司简称'], 公司全称=x.get('公司全称'), 备注=x.get('备注')) for x in ctx['companies']]
    if k == '账户':
        subs = {s for s in SUB_NAMES}
        out = []
        for a in ctx['accounts']:
            sub = a.get('子表名') if a.get('子表名') in subs else None
            out.append(dict(账户名称=a['账户名称'], 子表名=sub, 公司=a.get('公司'), 户别=a.get('户别'), 开户行=a.get('开户行'),
                            户号尾号=a.get('户号尾号'), 持卡人=a.get('持卡人'),
                            期初余额=0 if a.get('期初余额') is None else a.get('期初余额'),
                            状态=a.get('状态'), 停用日=a.get('停用日'), 别名=a.get('别名'), 备注=a.get('备注')))
        return out
    if k == '收费项':
        return [dict(收费项=x['收费项'], 属性=x.get('属性'), 大类=x.get('大类'), 别名=x.get('别名'), 说明=x.get('说明'))
                for x in ctx['items']]
    if k == '收付人':
        return [dict(收付人=x['收付人'], 类型=x.get('类型'), 别名=x.get('别名'), 备注=x.get('备注')) for x in ctx['parties']]
    if k == '规则':
        return [dict(关键词=x.get('关键词'), 查找位置=x.get('查找位置'), 方向=x.get('方向'), 适用范围=x.get('适用范围'),
                     收费项=x.get('收费项'), 收付人=x.get('收付人'), 规则号=x.get('规则号'), 说明=x.get('说明'))
                for x in ctx['rules']]
    return []


def _param_now(key):
    """① 参数右边「现在按 … 算」"""
    b = f'${BASE_PAR_VAL}${PAR_ROW[key]}'
    if key == '集团名称':
        return '=IF(P_集团名称="","（空着：标题写「集团」）","标题写「"&P_集团名称&"」")'
    if key == '管理年度':
        return (f'=IF(AND({raw(b)},NOT(AND(ISNUMBER({b}),{b}>=2000,{b}<=2100))),"※ 看不懂，按 "&P_管理年度&" 年算",'
                f'"现在按 "&P_管理年度&" 年"&IF({raw(b)},"","（最后一笔流水的年份）"))')
    if key == '开账日':
        d = 'YEAR(P_开账日)&"/"&MONTH(P_开账日)&"/"&DAY(P_开账日)'
        return f'=IF({dsel_bad(b)},"※ 看不懂，按 "&{d}&" 算","现在按 "&{d}&IF({raw(b)},"","（空着）"))'
    if key == '报表暂停':
        return '=IF(P_报表暂停=1,"七张报表暂停中（清空就恢复）","照常算")'
    return None


def build_base(ws, ctx):
    widths(ws, BASE_WIDTH)
    title(ws, '=IF(P_集团名称="","集团",P_集团名称)&" 基础资料"', 'N', BASE_TIP)
    # 第 3 行：跳到各块
    nav = [('① 参数', PAR_TITLE)] + [(t.split('（')[0].strip(), tr) for _k, t, tr, _h, _n in BASE_BLOCKS]
    for (txt, row), col in zip(nav, ['B', 'C', 'D', 'E', 'F', 'G']):
        c = put(ws, f'{col}3', txt, F_LINK, fill('FFF3EAFB'), align=AC)
        link(c, SH_BASE, f'A{row}')
    put(ws, 'A3', '跳到：', F_KPI_L, align=AC, border=False)
    home_link(ws, 'I3')
    ws.row_dimensions[3].height = 22
    # ① 参数
    section(ws, PAR_TITLE, 'A', 'N', '① 参数（淡黄格可以改）', group='基础')
    header(ws, PAR_HDR, [('A', '项目'), ('B', '值'), ('C', '说明'), ('L', '现在按')], '基础', height=24)
    ws.merge_cells(f'C{PAR_HDR}:K{PAR_HDR}')
    ws.merge_cells(f'L{PAR_HDR}:N{PAR_HDR}')
    pvals = dict(ctx.get('params') or {})
    for key, label, default, note in PARAMS:
        r = PAR_ROW[key]
        v = pvals.get(key)
        if key == '开账日' and v in (None, ''):
            v = dt.datetime.strptime(default, '%Y-%m-%d')           # 用户给的值（哪怕是文字）照写，公式 dsel 会认
        put(ws, f'A{r}', label, F_TXTB, fill('FFF3EAFB'), align=AC)
        fmt = {'管理年度': FMT_YEAR, '开账日': DATE}.get(key)
        cell = put(ws, f'B{r}', None, F_SEL, FILL_IN, fmt=fmt, align=AC)
        if v not in (None, ''):
            _set(ws, r, 2, v)
        put(ws, f'C{r}', note, F_NOTE, align=ALW)
        ws.merge_cells(f'C{r}:K{r}')
        put(ws, f'L{r}', _param_now(key), F_AUTO, FILL_AUTO, align=AL)
        ws.merge_cells(f'L{r}:N{r}')
        ws.row_dimensions[r].height = 30
    cf_flags(ws, f'L{PAR_ROW[PARAMS[0][0]]}:L{PAR_ROW[PARAMS[-1][0]]}', f'L{PAR_ROW[PARAMS[0][0]]}')
    dv_add(ws, f'B{PAR_ROW["管理年度"]}', type='whole', operator='between', formula1='2000', formula2='2100', stop=True,
           prompt='空着＝最后一笔流水的年份。1 月份想看上一年全年，填上一年（如 2026），看完清空',
           err_title='管理年度', err='填 4 位年份，比如 2026；空着＝最后一笔流水的年份')
    dv_add(ws, f'B{PAR_ROW["开账日"]}', type='date', operator='between', formula1='36526', formula2='73050', stop=True,
           prompt='本本从这一天开始（一年一本，填 1 月 1 日）。③的期初余额是这一天开始前的余额',
           err_title='开账日', err='要填真正的日期，像 2026/1/1')
    dv_fixed(ws, f'B{PAR_ROW["报表暂停"]}', '是', stop=False,
             prompt='一次录很多时填「是」：七张报表先不算，打字不卡；首页、子表、台帐照常算。录完记得清空')
    # ②～⑥ 各块
    for k, ttl, tr, h, cap in BASE_BLOCKS:
        cols = BASE_COLS[k]
        lastc = max(cols.values(), key=CI)
        section(ws, tr, 'A', lastc, ttl, group='基础')
        put(ws, f'{CL(CI(lastc) + 1)}{tr}', '↑ 回顶上', F_LINK, align=AC, border=False)
        link(ws[f'{CL(CI(lastc) + 1)}{tr}'], SH_BASE, 'A1')
        header(ws, h, [(col, BASE_HEAD[k][f]) for f, col in cols.items()], '基础', height=44)
        r0, r1 = _block_rows(k)
        rows = _preset_rows(k, ctx)
        if len(rows) > cap:
            print(f'   ⚠ s_input：{k} {len(rows)} 个超过容量 {cap}，多的不写')
        st = _BASE_STYLE[k]
        for i in range(cap):
            r = r0 + i
            x = rows[i] if i < len(rows) else {}
            for f, col in cols.items():
                if f == '序号':
                    c = ws.cell(row=r, column=CI(col), value=f'=IF(TRIM($B{r}&"")="","",ROW()-ROW($A${h}))')
                    _apply(c, _ST_SEQ)
                elif f == '核对':
                    c = ws.cell(row=r, column=CI(col), value=_check_formula(k, r))
                    _apply(c, _ST_CHK)
                else:
                    _set(ws, r, CI(col), x.get(f), st[f])
        cf_flags(ws, f'{cols["核对"]}{r0}:{cols["核对"]}{r1}', f'{cols["核对"]}{r0}')
        # 有效性
        rg = lambda f: f'{cols[f]}{r0}:{cols[f]}{r1}'
        if k == '账户':
            dv_list(ws, rg('公司'), '公司列表', stop=True, prompt='选②里的公司简称（新公司先去②加）',
                    error='公司要先在②里加，再回来选')
            dv_fixed(ws, rg('户别'), '户别', prompt='对公基本户/对公一般户/收单户 算「对公」；私户、现金单列（首页小计用）')
            dv_add(ws, rg('期初余额'), type='decimal', operator='between', formula1='-999999999999', formula2='999999999999',
                   stop=True, prompt='开账日那天开始前的余额。年中新开的账户填 0，开户那天的余额在台帐录一笔「账户转入」',
                   err_title='期初余额', err='要填数字（没有就填 0），不要带文字、逗号以外的符号')
            dv_fixed(ws, rg('状态'), '状态', prompt='空＝在用。不用的账户改「销户」并填停用日，别删行（删了后面的账户号会往前挪）')
            dv_add(ws, rg('停用日'), type='date', operator='between', formula1='36526', formula2='73050', stop=True,
                   prompt='暂停/销户从哪天起。这天以前的流水不提示「暂停户还在记」',
                   err_title='停用日', err='要填真正的日期，像 2026/3/21')
            dv_add(ws, rg('子表名'), prompt='生成时写死的：子表靠它认账户，别改。新开的卡用备用子表（在它 D3 选账户），这里空着')
            dv_add(ws, rg('别名'), prompt='原表名、旧名字，用「、」隔开。台帐里写这些名字也认成这个账户')
        elif k == '收费项':
            dv_fixed(ws, rg('属性'), '属性',
                     prompt='收入/支出算经营；投资建设、往来、待分类不算经营；内部转账不算收支；按方向＝收入栏算收入、支出栏算支出')
            dv_add(ws, rg('大类'), prompt=f'报表里按大类分组小计（最多 {N_GROUP} 种）')
            dv_add(ws, rg('别名'), prompt='以前的写法，用「、」隔开。台帐里写这些也认成这个收费项')
        elif k == '收付人':
            dv_fixed(ws, rg('类型'), '收付人类型', prompt='收支报表按类型合计用')
            dv_add(ws, rg('别名'), prompt='以前的写法，用「、」隔开。台帐里写这些也认成这个收付人')
        elif k == '规则':
            dv_add(ws, rg('关键词'), prompt='摘要（或收付人）里包含这几个字就算命中。从上往下，先命中的那条算；越具体的放越上面')
            dv_fixed(ws, rg('查找位置'), '查找位置', prompt='在摘要里找还是在收付人里找；空＝两处都找')
            dv_fixed(ws, rg('方向'), '方向', prompt='收＝只认收入的笔，支＝只认支出的笔，都（或空）＝都认')
            dv_add(ws, rg('适用范围'), prompt='只对某家公司（②的简称）或某个账户（③的名称）起作用；空＝全部')
            dv_list(ws, rg('收费项'), '收费项列表', stop=False, prompt='命中后认成这个收费项（不填＝这条只认收付人）')
            dv_list(ws, rg('收付人'), '收付人列表', stop=False, prompt='命中后认成这个收付人（不填＝这条只认收费项）')
    put(ws, f'A{BASE_END_ROW}', BASE_END_TEXT, F_NOTE, align=AL, border=False)
    ws.freeze_panes = 'A4'
    print_setup(ws, None, landscape=True)


def build(wb, ctx):
    build_ledger(wb[SH_LEDGER], ctx)
    build_base(wb[SH_BASE], ctx)
