# -*- coding: utf-8 -*-
"""报表二（s_rep2）：【收支分类报表】【收支分项汇总】【收支报表】【任意时间段收支报表】。设计见 设计.md §5.5～§5.8。

口径（设计 §1.3 经营口径）：经营收入＝属性「收入」行 Σ(收－支)；经营支出＝属性「支出」行 Σ(支－收)；投资建设＝Σ(支－收)（正数＝花出去）；
  往来净额、待分类净额、内部净额＝Σ(收－支)。恒等式：经营净额－投资建设＋往来净额＋待分类净额＋内部净额＝资金净增加（＝Σ 收－支）。
  收支分类报表选了公司、收支分项汇总的公司列都按「归属公司」（流_归年属月键 / 流_项归年属月键）；任意时间段的「各账户」是账户口径。
写法：
  - 只引用定义名称（_流、_表、_参、块_）和本表格子；扫 _流 一律配窗口（每个月一对窗口放本表隐藏格），条件照写全（窗口只是少扫）。
  - 收支分类、分项汇总的行＝_表 分类_候选键 第 k 小（段首、大类小计、收费项、认不出）；段首、小计只加「它下面」的行：
    SUMIFS(本列下一行:本列表尾下一行, …)，引用一律往下，Excel 静态判环也不成环（不能 SUMIFS 整列：会把自己圈进去）。
  - 每块「其余 / 未列示 / 公司没认出」＝总额－明细（补差），保证每块合计＝总额。
  - 第 2 行起每个公式都套 layout.rpt()（报表暂停）；第 1 行标题暂停时红字。
私有帮助格（不进 layout.NAMES，只本表用）：
  收支分类报表：R1..R4 标签、Z2..AA13（1～12 月窗口）、I3:L3（本年经营净额）、第 416 行留空（小计的区域尾）；
  收支分项汇总：AI 列「公司没认出」（＝合计－各公司，合计里含）、AK2..AK4（起月、止月、年份格生效值）、AP1..AP4 标签、第 415 行留空；
  收支报表：Q 列（块 1～3 每行的收付人号 / 收费项号）、Y 列（块 2 每行本年笔数）、Y1..Y3 标签、AA2..AB13（1～12 月窗口）、
    第 249 行（块 4「没归到 9 类的」）；
  任意时间段：X1..X3 标签、Z/AA 第 8..47 行（账户号、期间笔数）、AB..AE 第 54 行起（四块每行的收费项号 / 收付人号）、K3:L4（资金净增加、核对）、
    第 5 行（选择格「按 … 算」）。"""
import re
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from layout import (
    SH_CLS, SH_ITEMCO, SH_IE, SH_SPAN, N_CO, N_CLS, N_PARTY, K_DAY, ATTR_NO, PARTY_TYPES, DV_FIXED,
    CLS, CLS_SEL, CLS_EFF, ITEMCO, ITEMCO_SEL, ITEMCO_EFF, IE, IE_SEL, IE_EFF, IE_TOP_IN, IE_TOP_OUT, IE_ITEMS,
    SPAN_SEL, SPAN_KPI, SPAN_EFF, SPAN_ACC, SPAN_BLK, SPAN_RANK,
    rpt, dsel_bad, yr_eff, yr_bad, ym_eff, ym_of, raw, win_first, win_last, win, k_aypm, k_igypm, name2no)
from common import (
    YH, F_HDR, F_TXT, F_TXTB, F_AUTO, F_AUTOB, F_HELP, F_RED, F_KPI_L, F_KPI_V, F_NOTE, FILL_KPI, FILL_SUB, FILL_SECT,
    FILL_NONE, BD, AC, ACW, AL, AR, ACC, INT, PCT, put, title, header, section, widths, hide, selector, home_link,
    cf_flags, cf_neg_red, cf_bg, fill)

I9 = ATTR_NO['内部转账']
A_IN, A_OUT, A_CAP, A_FLOW, A_TBD = (ATTR_NO[k] for k in ('收入', '支出', '投资建设', '往来', '待分类'))
PS = [A_IN, A_OUT, A_CAP, A_FLOW, A_TBD, I9]                        # 段的顺序（＝分类_候选键 从小到大）
SEC = {A_IN: '一、收入（经营收入）', A_OUT: '二、支出（经营支出）', A_CAP: '三、投资建设（正数＝花出去）',
       A_FLOW: '四、往来（正数＝净流入）', A_TBD: '五、待分类（请定性质）', I9: '六、内部转账（全集团应为 0）'}
ATXT = {A_IN: '收入', A_OUT: '支出', A_CAP: '投资建设', A_FLOW: '往来', A_TBD: '待分类', I9: '内部转账'}
UNKNOWN_TXT = '　　（收费项没填或认不出）'
GRP99 = '（大类空或超过 40 个）'
ASHRINK = Alignment(horizontal='left', vertical='center', shrink_to_fit=True)
ACSHRINK = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
F_SUBT = Font(name=YH, sz=10, bold=True, color='FF1F3864')
TYPE_SEC, TYPE_GRP, TYPE_ITEM, TYPE_UNK = 1, 2, 3, 4                # 收支分类/分项汇总每行的类型（帮助列）


def R(f):
    return rpt(f)


def ab(c):
    m = re.fullmatch(r'([A-Z]+)(\d+)', c)
    return f'${m.group(1)}${m.group(2)}'


def d_txt(x):
    return f'YEAR({x})&"/"&MONTH({x})&"/"&DAY({x})'


def ym_txt(x):
    return f'INT({x}/100)&"年"&MOD({x},100)&"月"'


def choose_p(p, d):
    """属性号 p（1..5、9）→ d[p]；不是这几个号＝""（CHOOSE 第 6～8 个参数空着）"""
    arg = ','.join(f'"{d.get(i, "")}"' for i in range(1, I9 + 1))
    return f'IF(OR({p}<1,{p}>{I9}),"",CHOOSE({p},{arg}))'


def sign(p):
    """支出、投资建设取正数（Σ 支－收）；其余 Σ 收－支"""
    return f'IF(OR({p}={A_OUT},{p}={A_CAP}),-1,1)'


def pause_title(ws, f):
    ws['A1'] = f'=IF(P_报表暂停=1,"报表暂停中：基础资料①「报表暂停」清空就恢复",{f[1:] if f.startswith("=") else f})'
    ws.conditional_formatting.add('A1', FormulaRule(formula=['P_报表暂停=1'], font=Font(name=YH, sz=15, bold=True, color='FFC00000'),
                                                    fill=cf_bg('FFFFD7D7')))


def group_name():
    return 'IF(P_集团名称="","集团",P_集团名称)'


def money(ws, coord, f, bold=False, fill_=None, font=None):
    return put(ws, coord, f, font or (F_AUTOB if bold else F_AUTO), fill_ or FILL_NONE, fmt=ACC, align=AR)


def txt(ws, coord, f, bold=False, fill_=None, align=None, font=None):
    return put(ws, coord, f, font or (F_TXTB if bold else F_TXT), fill_ or FILL_NONE, align=align or AL)


def pct(ws, coord, f, font=None, fill_=None):
    return put(ws, coord, f, font or F_AUTO, fill_ or FILL_NONE, fmt=PCT, align=AR)


def helper(ws, coord, f):
    c = ws[coord]
    c.value = f
    c.font = F_HELP
    return c


def merge_put(ws, rng, f, font, fill_, fmt=None, align=AC):
    ws.merge_cells(rng)
    c1 = rng.split(':')[0]
    put(ws, c1, f, font, fill_, fmt=fmt, align=align)
    m = re.fullmatch(r'([A-Z]+)(\d+):([A-Z]+)(\d+)', rng)
    for ci in range(CI(m.group(1)), CI(m.group(3)) + 1):
        for r in range(int(m.group(2)), int(m.group(4)) + 1):
            ws.cell(row=r, column=ci).border = BD


def by_year(sel, eff):
    return R(f'IF({yr_bad(sel)},"※ 看不懂，按 "&{eff}&" 年算",IF({eff}<P_开账年,"※ 本本从 "&P_开账年&" 年开始，"&{eff}&" 年的看上一本",'
             f'"按 "&{eff}&" 年算"&IF({raw(sel)},"","（管理年度）")))')


def grp_label(v):
    return f'IF({v}=99,"{GRP99}",INDEX(大类_名称,{v}))'


def money_txt(x):
    return f'TEXT(ABS({x}),"#,##0.00")'


# ─────────── 收支分类 / 分项汇总共用：候选键行（段首、大类小计、收费项、认不出） ───────────
def cand_rows(ws, r1, r2, H, C):
    """帮助列 键/类型/属性号/大类号/项号 ＋ 显示列 项目/属性/大类（第 r1..r2 行）。H＝帮助列字母，C＝显示列字母"""
    S, T, U, V, W = (H[k] for k in ('键', '类型', '属性号', '大类号', '项号'))
    for r in range(r1, r2 + 1):
        s, t, u, v, w = (f'${x}{r}' for x in (S, T, U, V, W))
        helper(ws, f'{S}{r}', R(f'IFERROR(SMALL(分类_候选键,ROWS(${S}${r1}:{S}{r})),0)'))
        helper(ws, f'{T}{r}', R(f'IF({s}=0,0,IF(MOD({s},1000000)=0,{TYPE_SEC},IF(MOD({s},1000000)=999999,{TYPE_UNK},'
                                f'IF(MOD({s},1000)=0,{TYPE_GRP},{TYPE_ITEM}))))'))
        helper(ws, f'{U}{r}', R(f'IF({s}=0,0,INT({s}/1000000))'))
        helper(ws, f'{V}{r}', R(f'IF(OR({t}={TYPE_GRP},{t}={TYPE_ITEM}),INT(MOD({s},1000000)/1000),0)'))
        helper(ws, f'{W}{r}', R(f'IF({t}={TYPE_ITEM},MOD({s},1000),0)'))
        name = (f'IF({t}=0,"",IF({t}={TYPE_SEC},{choose_p(u, SEC)},IF({t}={TYPE_GRP},"　"&{grp_label(v)}&" 小计",'
                f'IF({t}={TYPE_ITEM},"　　"&INDEX(项_名称,{w}),"{UNKNOWN_TXT}"))))')
        txt(ws, f'{C["项目"]}{r}', R(name), align=ASHRINK)
        txt(ws, f'{C["属性"]}{r}', R(f'IF(OR({t}=0,{t}={TYPE_SEC}),"",{choose_p(u, ATXT)}&IF({t}={TYPE_ITEM},'
                                    f'IF(INDEX(项_属性号,{w})=0,"（按方向）",""),""))'), align=ACSHRINK)
        txt(ws, f'{C["大类"]}{r}', R(f'IF(OR({t}={TYPE_GRP},{t}={TYPE_ITEM}),{grp_label(v)},"")'), align=ASHRINK)


def sub_total(col, r, H, end):
    """段首、大类小计：只加下面的行（本列 r+1..end；end＝表尾下一行，留空）"""
    T, U, V = H['类型'], H['属性号'], H['大类号']
    rng = lambda c: f'{c}{r + 1}:{c}${end}'
    rngA = lambda c: f'${c}{r + 1}:${c}${end}'
    return (f'IF(${T}{r}={TYPE_GRP},SUMIFS({rng(col)},{rngA(U)},${U}{r},{rngA(V)},${V}{r},{rngA(T)},{TYPE_ITEM}),'
            f'SUMIFS({rng(col)},{rngA(U)},${U}{r},{rngA(T)},">={TYPE_ITEM}"))')


def cand_cf(ws, rng, tcell):
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{tcell}={TYPE_SEC}'], font=Font(name=YH, bold=True, color='FF375623'),
                                                   fill=cf_bg('FFE2EFDA')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{tcell}={TYPE_GRP}'], font=Font(name=YH, bold=True, color='FF1F3864'),
                                                   fill=cf_bg('FFDDEBF7')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{tcell}={TYPE_UNK}'], font=Font(name=YH, italic=True, color='FF8B5E00')))


# ═════════════════════════════ 收支分类报表 ═════════════════════════════
CLS_MWIN = dict(窗首='Z', 窗末='AA')                       # 第 m+1 行＝m 月的窗口（私有）


def build_cls(ws):
    C, H = CLS['列'], CLS['帮助']
    E = {k: ab(v) for k, v in CLS_EFF.items()}
    S = {k: ab(v) for k, v in CLS_SEL.items()}
    Y, G = E['年份'], E['公司号']
    mcol = [CL(CI(C['月1']) + m - 1) for m in range(1, 13)]
    assert CI(C['全年']) == CI(C['月1']) + 12 and CI(C['占比']) == CI(C['全年']) + 1
    r1, r2 = CLS['首行'], CLS['末行']
    end = r2 + 1                                             # 小计区域尾（留空）
    widths(ws, {C['项目']: 30, C['属性']: 11, C['大类']: 14, C['全年']: 15, C['占比']: 9})
    for c in mcol:
        ws.column_dimensions[c].width = 12.5
    tip = ('经营口径：收入、支出只算收费项属性是「收入」「支出」的（收入类记在支出栏的算冲减，比如退租金）；投资建设（正数＝花出去）、'
           '往来（借款、还款、押金…）、待分类（收费项待定、没填或认不出的）、内部转账单列，不算经营。上面是汇总，下面按属性 → 大类 → 收费项展开，'
           '每段最后一行是「收费项没填或认不出」的。黄格：年份空着＝管理年度；公司空着＝全集团，选了公司按「归属公司」算'
           '（私户替别的公司收付时台帐填的归属公司，没填＝账户所属公司），这时内部净额不为 0 是正常的。名单来自【基础资料】④，'
           '要细分就去那边加收费项、定大类。全部自动算，不用填。')
    title(ws, '=""', C['占比'], tip, h2=58)
    co_nm = f'IF({G}>0,INDEX(公司_名称,{G}),{group_name()})'
    pause_title(ws, f'{co_nm}&" "&{Y}&" 年收支分类（经营口径）"')
    # —— 生效值（隐藏）——
    f = dict(年份=yr_eff(S['年份']), 公司号=name2no(S['公司'], '公司'),
             窗首=win_first(f'{Y}*100+1'), 窗末=win_last(f'{Y}*100+12', E['窗首']))
    for k, v in f.items():
        helper(ws, CLS_EFF[k], R(v))
        helper(ws, f'R{CLS_EFF[k][1:]}', k)
    W1, W2 = CLS_MWIN['窗首'], CLS_MWIN['窗末']
    helper(ws, f'{W1}1', '月窗首')
    helper(ws, f'{W2}1', '月窗末')
    for m in range(1, 13):
        helper(ws, f'{W1}{m + 1}', R(win_first(f'{Y}*100+{m}')))
        helper(ws, f'{W2}{m + 1}', R(win_last(f'{Y}*100+{m}', f'{W1}{m + 1}')))
    wm = lambda nm, m: win(nm, f'${W1}${m + 1}', f'${W2}${m + 1}')
    wy = lambda nm: win(nm, E['窗首'], E['窗末'])
    # —— 第 3 行 ——
    selector(ws, 'A3', '年份', CLS_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    txt(ws, 'C3', by_year(S['年份'], Y), font=F_NOTE, align=ASHRINK)
    selector(ws, 'D3', '公司', CLS_SEL['公司'], dv_formula='公司列表', prompt='选一家公司（按归属公司算）；空着＝全集团')
    sel_c = f'TRIM({S["公司"]}&"")'
    merge_put(ws, 'F3:H3', R(f'IF({sel_c}="","按 全集团 算",IF({G}=0,"※ 「"&{sel_c}&"」不在基础资料②，按全集团算",'
                             f'"按 "&INDEX(公司_名称,{G})&" 算（归属公司）"))'), F_NOTE, FILL_NONE, align=ASHRINK)
    merge_put(ws, 'I3:J3', '本年经营净额', F_KPI_L, FILL_KPI, align=ACW)
    merge_put(ws, 'K3:L3', R(f'{C["全年"]}{CLS["汇总"]["经营净额"]}'), F_KPI_V, FILL_KPI, fmt=ACC)
    cf_flags(ws, 'F3', 'F3')
    home_link(ws, 'P3')
    ws.row_dimensions[3].height = 24
    # —— 表头 ——
    hdr = [(C['项目'], '项目'), (C['属性'], '属性'), (C['大类'], '大类')] + [(c, f'{m}月') for m, c in enumerate(mcol, 1)] + \
          [(C['全年'], '全年合计'), (C['占比'], '占比')]
    header(ws, CLS['表头行'], hdr, '报表', height=24)
    header(ws, CLS['表头行2'], hdr, '报表', height=24)
    # —— 汇总（第 5～13 行）——
    SR = CLS['汇总']
    note = dict(经营收入='属性「收入」的行：收入栏－支出栏', 经营支出='属性「支出」的行：支出栏－收入栏', 经营净额='经营收入－经营支出',
                投资建设='买资产、建项目（正数＝花出去）', 往来净额='借款、还款、押金、投资款…（流入－流出）',
                待分类='收费项待定、没填或认不出（请定性质）', 内部净额='本册账户之间互转（全集团应为 0）',
                净增加='＝经营净额－投资建设＋往来＋待分类＋内部', 核对='资金净增加－台帐同期收－支（应为 0）')
    label = dict(经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设', 往来净额='往来净额',
                 待分类='待分类净额', 内部净额='内部转账净额', 净增加='资金净增加', 核对='核对')

    def sp(p, m):
        """m 月属性 p 的 Σ净（全集团按月码＋属性；选了公司按归属年属月键）"""
        return (f'IF({G}=0,SUMIFS({wm("流_净", m)},{wm("流_月码", m)},{Y}*100+{m},{wm("流_属性号", m)},{p}),'
                f'SUMIF({wm("流_归年属月键", m)},{k_aypm(G, Y, p, m)},{wm("流_净", m)}))')
    for k, r in SR.items():
        bold = k in ('经营净额', '净增加')
        txt(ws, f'{C["项目"]}{r}', label[k], bold=True, fill_=FILL_KPI if bold else FILL_SUB)
        merge_put(ws, f'{C["属性"]}{r}:{C["大类"]}{r}', note[k], F_NOTE, FILL_NONE, align=ASHRINK)
        for m, c in enumerate(mcol, 1):
            v = dict(经营收入=sp(A_IN, m), 经营支出=f'-{sp(A_OUT, m)}', 经营净额=f'{c}{SR["经营收入"]}-{c}{SR["经营支出"]}',
                     投资建设=f'-{sp(A_CAP, m)}', 往来净额=sp(A_FLOW, m), 待分类=sp(A_TBD, m), 内部净额=sp(I9, m),
                     净增加=(f'{c}{SR["经营净额"]}-{c}{SR["投资建设"]}+{c}{SR["往来净额"]}+{c}{SR["待分类"]}'
                          f'+{c}{SR["内部净额"]}'),
                     核对=(f'{c}{SR["净增加"]}-IF({G}=0,SUMIF({wm("流_月码", m)},{Y}*100+{m},{wm("流_净", m)}),'
                         f'SUMIFS({wm("流_净", m)},{wm("流_月码", m)},{Y}*100+{m},{wm("流_归属号", m)},{G}))'))[k]
            money(ws, f'{c}{r}', R(f'ROUND({v},2)'), bold=bold)
        if k == '核对':
            yv = (f'{C["全年"]}{SR["净增加"]}-IF({G}=0,SUMIFS({wy("流_净")},{wy("流_月码")},">="&({Y}*100+1),{wy("流_月码")},"<="&({Y}*100+12)),'
                  f'SUMIFS({wy("流_净")},{wy("流_月码")},">="&({Y}*100+1),{wy("流_月码")},"<="&({Y}*100+12),{wy("流_归属号")},{G}))')
            money(ws, f'{C["全年"]}{r}', R(f'ROUND({yv},2)'), bold=True)
        else:
            money(ws, f'{C["全年"]}{r}', R(f'ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2)'), bold=True,
                  fill_=FILL_KPI if bold else None, font=F_RED if bold else None)
        put(ws, f'{C["占比"]}{r}', None, F_AUTO, FILL_NONE)
    for k in ('经营净额', '待分类', '净增加'):
        r = SR[k]
        cf_neg_red(ws, f'{mcol[0]}{r}:{C["全年"]}{r}', f'{mcol[0]}{r}')
    ws.conditional_formatting.add(f'{mcol[0]}{SR["核对"]}:{C["全年"]}{SR["核对"]}',
                                  FormulaRule(formula=[f'AND(ISNUMBER({mcol[0]}{SR["核对"]}),ROUND({mcol[0]}{SR["核对"]},2)<>0)'],
                                              font=Font(name=YH, bold=True, color='FFC00000'), fill=cf_bg('FFFFD7D7')))
    # —— 明细（第 16～415 行）——
    cand_rows(ws, r1, r2, H, C)
    T, U, Wc, X = H['类型'], H['属性号'], H['项号'], H['笔数']
    wy_i, wy_ig = wy('流_项年属月键'), wy('流_项归年属月键')
    for r in range(r1, r2 + 1):
        t, u, w = f'${T}{r}', f'${U}{r}', f'${Wc}{r}'
        helper(ws, f'{X}{r}', R(f'IF({t}<{TYPE_ITEM},0,IF({G}=0,COUNTIFS({wy_i},">="&{k_aypm(w, Y, u, 1)},{wy_i},"<="&{k_aypm(w, Y, u, 12)}),'
                                f'COUNTIFS({wy_ig},">="&{k_igypm(w, G, Y, u, 1)},{wy_ig},"<="&{k_igypm(w, G, Y, u, 12)})))'))
        for m, c in enumerate(mcol, 1):
            one = (f'IF({G}=0,SUMIF({wm("流_项年属月键", m)},{k_aypm(w, Y, u, m)},{wm("流_净", m)}),'
                   f'SUMIF({wm("流_项归年属月键", m)},{k_igypm(w, G, Y, u, m)},{wm("流_净", m)}))')
            money(ws, f'{c}{r}', R(f'IF({t}=0,"",IF({t}>={TYPE_ITEM},IF(${X}{r}=0,0,ROUND({sign(u)}*{one},2)),'
                                   f'{sub_total(c, r, H, end)}))'))
        money(ws, f'{C["全年"]}{r}', R(f'IF({t}=0,"",ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2))'), bold=True)
        pin, pout = f'${C["全年"]}${SR["经营收入"]}', f'${C["全年"]}${SR["经营支出"]}'
        pct(ws, f'{C["占比"]}{r}', R(f'IF(OR({t}=0,AND({u}<>{A_IN},{u}<>{A_OUT})),"",IF({u}={A_IN},IF(ROUND({pin},2)=0,"",'
                                     f'{C["全年"]}{r}/{pin}),IF(ROUND({pout},2)=0,"",{C["全年"]}{r}/{pout})))'))
    cand_cf(ws, f'A{r1}:{C["占比"]}{r2}', f'${T}{r1}')
    # —— 提示（第 14 行）——
    tr = CLS['提示行']
    P = C['全年']
    sec_sum = {A_IN: SR['经营收入'], A_OUT: SR['经营支出'], A_CAP: SR['投资建设'], A_FLOW: SR['往来净额'], A_TBD: SR['待分类'],
               I9: SR['内部净额']}
    mis = '+'.join(f'(ROUND(SUMIFS(${P}${r1}:${P}${r2},${T}${r1}:${T}${r2},{TYPE_SEC},${U}${r1}:${U}${r2},{p})-${P}${rr},2)<>0)'
                   for p, rr in sec_sum.items())
    tbd, intl, chk = f'${P}${SR["待分类"]}', f'${P}${SR["内部净额"]}', f'${P}${SR["核对"]}'
    tipf = (f'IF(COUNT(分类_候选键)>{N_CLS},"※ 收费项、大类太多：本表只放得下 {N_CLS} 行，后面的没列出来（下面各段会比上面汇总少）",'
            f'IF(({mis})>0,"※ 下面明细跟上面汇总对不上 "&({mis})&" 段：请告诉我们",'
            f'IF(ROUND({chk},2)<>0,"※ 资金净增加跟台帐对不上：请告诉我们",'
            f'IF(ROUND({tbd},2)<>0,"△ 待分类净额 "&IF({tbd}<0,"-","")&{money_txt(tbd)}&'
            f'"：收费项待定、没填或认不出，去【资金台帐】选收费项（或在基础资料④定属性）",'
            f'IF(AND({G}=0,ROUND({intl},2)<>0),"△ 内部转账全年差 "&{money_txt(intl)}&'
            f'"：漏录一头、金额或日期对不上，看核对表清单","✓ 下面明细合计跟上面汇总一致"&IF({G}>0,"（选了公司：内部净额不为 0 是正常的）",""))))))')
    merge_put(ws, f'A{tr}:{C["占比"]}{tr}', R(tipf), F_TXT, FILL_NONE, align=AL)
    cf_flags(ws, f'A{tr}', f'A{tr}')
    ws.row_dimensions[tr].height = 22
    hide(ws, 'R', *H.values(), W1, W2)
    ws.freeze_panes = f'{mcol[0]}{r1}'


# ═════════════════════════════ 收支分项汇总（收费项 × 公司）═════════════════════════════
IC_PRIV = dict(起月='AK2', 止月='AK3', 年份格='AK4')        # 私有生效值；标签在 AP 列
IC_UNL = 'AI'                                              # 「公司没认出」（＝合计－各公司；合计里含）


def build_itemco(ws, n_co):
    C, H = ITEMCO['列'], ITEMCO['帮助']
    E = {k: ab(v) for k, v in {**ITEMCO_EFF, **IC_PRIV}.items()}
    S = {k: ab(v) for k, v in ITEMCO_SEL.items()}
    Y, M1, M2, YM, FY, YG = E['年份'], E['起月'], E['止月'], E['月码'], E['全年'], E['年份格']
    c0 = CI(C['公司1'])
    ccol = [CL(c0 + k - 1) for k in range(1, N_CO + 1)]
    TOT = C['合计']
    assert CI(TOT) == c0 + N_CO and CI(IC_UNL) == CI(TOT) + 1 and CI(H['键']) == CI(IC_UNL) + 1
    r1, r2 = ITEMCO['首行'], ITEMCO['末行']
    end = r2 + 1
    widths(ws, {C['项目']: 28, C['属性']: 11, C['大类']: 14, TOT: 15, IC_UNL: 12})
    for c in ccol:
        ws.column_dimensions[c].width = 12.5
    tip = ('像原来日记账的【汇总】：每个收费项各家公司收了多少、花了多少。按「归属公司」算（私户替别的公司收付时台帐填的归属公司，没填＝账户所属公司）；'
           '经营口径：支出、投资建设取正数（花出去），其余是收－支。上面是各公司的汇总，下面按属性 → 大类 → 收费项展开。'
           '黄格：年份空着＝管理年度；月份空着＝首页截止日那个月（填 1～12＝上面年份的那个月）；「看全年」选「是」看一整年（要扫一年，慢一点）。'
           '合计列跟【收支分类报表】同期全集团的数一样；「公司没认出」是账户的公司不在基础资料②的（算在合计里，不在哪家公司）。'
           '公司超过现在的家数以后的列隐藏了：加了公司就取消隐藏。全部自动算，不用填。')
    title(ws, '=""', TOT, tip, h2=58)
    pause_title(ws, f'{group_name()}&" "&{Y}&" 年 "&IF({FY}=1,"全年",MOD({YM},100)&" 月")&" 收支分项汇总（收费项 × 公司，按归属公司）"')
    # —— 生效值（隐藏）——
    hs = f'TRIM({S["全年"]}&"")'
    f = [('年份格', yr_eff(S['年份'])),
         ('全年', f'IF(OR({hs}="是",{hs}="全年",{hs}="1"),1,0)'),
         ('月码', f'IF({FY}=1,0,{ym_eff(S["月份"], YG, f"IF({YG}=P_截止年,P_截止月码,{YG}*100+12)")})'),
         ('年份', f'IF({FY}=1,{E["年份格"]},INT({YM}/100))'),
         ('起月', f'IF({FY}=1,1,MOD({YM},100))'), ('止月', f'IF({FY}=1,12,MOD({YM},100))'),
         ('窗首', win_first(f'{Y}*100+{M1}')), ('窗末', win_last(f'{Y}*100+{M2}', E['窗首']))]
    cells = {**ITEMCO_EFF, **IC_PRIV}
    for k, v in f:
        helper(ws, cells[k], R(v))
    for i, k in enumerate(('年份', '月码', '窗首', '窗末'), 1):
        helper(ws, f'AP{i}', f'{k}（{cells[k]}）')
    helper(ws, 'AQ1', '全年（AK1）、起月（AK2）、止月（AK3）、年份格（AK4）')
    wy = lambda nm: win(nm, E['窗首'], E['窗末'])
    # —— 第 3 行 ——
    selector(ws, 'A3', '年份', ITEMCO_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    txt(ws, 'C3', by_year(S['年份'], E['年份格']), font=F_NOTE, align=ASHRINK)
    selector(ws, 'D3', '月份', ITEMCO_SEL['月份'], kind='month',
             prompt='填 1～12＝上面年份的那个月；也可以填 202609 或那个月的某一天；空着＝首页截止日那个月')
    bad_m = f'AND({raw(S["月份"])},{ym_eff(S["月份"], E["年份格"], 0)}=0)'
    txt(ws, 'F3', R(f'IF({FY}=1,"看全年（月份不管用）",IF({bad_m},"※ 看不懂，按 "&{ym_txt(YM)}&" 算","按 "&{ym_txt(YM)}&" 算"&'
                    f'IF({raw(S["月份"])},"","（截止日那个月）")))'), font=F_NOTE, align=ASHRINK)
    selector(ws, 'G3', '看全年', ITEMCO_SEL['全年'], dv_formula='"' + DV_FIXED['是'] + '"',
             prompt='选「是」看一整年（扫一年，慢一点）；空着＝只看上面那个月')
    txt(ws, 'I3', R(f'IF({FY}=1,"按 "&{Y}&" 年全年算",IF({raw(S["全年"])},"※ 只认「是」，按一个月算","空着＝只看一个月"))'),
        font=F_NOTE, align=ASHRINK)
    cf_flags(ws, 'C3:I3', 'C3')
    home_link(ws, 'K3')
    ws.row_dimensions[3].height = 24
    # —— 表头（第 4、14 行）——
    hr, hr2 = ITEMCO['表头行'], ITEMCO['表头行2']
    hdr = [(C['项目'], '项目'), (C['属性'], '属性'), (C['大类'], '大类'), (TOT, '合计'), (IC_UNL, '公司没认出\n（合计里含）')]
    header(ws, hr, hdr, '报表', height=34)
    header(ws, hr2, hdr, '报表', height=34)
    for k, c in enumerate(ccol, 1):
        put(ws, f'{c}{hr}', R(f'IF(INDEX(公司_有效,{k})=1,INDEX(公司_名称,{k}),"")'), F_HDR, fill('FF30561E'), align=ACW)
        put(ws, f'{c}{hr2}', R(f'{c}{hr}'), F_HDR, fill('FF30561E'), align=ACW)
    for k, t in (('键', '键'), ('类型', '类型'), ('属性号', '属性号'), ('大类号', '大类号'), ('项号', '项号'), ('笔数', '笔数')):
        helper(ws, f'{H[k]}{hr2}', t)
    # —— 汇总（第 5～12 行）——
    SR = ITEMCO['汇总']
    label = dict(经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设（正数＝花出去）', 往来净额='往来净额',
                 待分类='待分类净额', 内部净额='内部转账净额', 净增加='资金净增加')
    wg, wn, wmm, wp = wy('流_归年属月键'), wy('流_净'), wy('流_月码'), wy('流_属性号')
    sg = lambda k, p: f'SUMIFS({wn},{wg},">="&{k_aypm(k, Y, p, M1)},{wg},"<="&{k_aypm(k, Y, p, M2)})'
    sa = lambda p: f'SUMIFS({wn},{wmm},">="&({Y}*100+{M1}),{wmm},"<="&({Y}*100+{M2}),{wp},{p})'

    def vals(c, sp):
        return dict(经营收入=sp(A_IN), 经营支出=f'-{sp(A_OUT)}', 经营净额=f'{c}{SR["经营收入"]}-{c}{SR["经营支出"]}',
                    投资建设=f'-{sp(A_CAP)}', 往来净额=sp(A_FLOW), 待分类=sp(A_TBD), 内部净额=sp(I9),
                    净增加=f'{c}{SR["经营净额"]}-{c}{SR["投资建设"]}+{c}{SR["往来净额"]}+{c}{SR["待分类"]}+{c}{SR["内部净额"]}')
    for k, r in SR.items():
        bold = k in ('经营净额', '净增加')
        merge_put(ws, f'{C["项目"]}{r}:{C["大类"]}{r}', label[k], F_TXTB, FILL_KPI if bold else FILL_SUB, align=AL)
        for kk, c in enumerate(ccol, 1):
            v = vals(c, lambda p, kk=kk: sg(kk, p))[k]
            money(ws, f'{c}{r}', R(f'IF({c}${hr}="","",ROUND({v},2))'), bold=bold)
        money(ws, f'{TOT}{r}', R(f'ROUND({vals(TOT, sa)[k]},2)'), bold=True, fill_=FILL_KPI, font=F_RED)
        money(ws, f'{IC_UNL}{r}', R(f'ROUND({TOT}{r}-SUM({ccol[0]}{r}:{ccol[-1]}{r}),2)'))
    for k in ('经营净额', '待分类', '净增加'):
        cf_neg_red(ws, f'{ccol[0]}{SR[k]}:{IC_UNL}{SR[k]}', f'{ccol[0]}{SR[k]}')
    # —— 明细（第 15～414 行）——
    cand_rows(ws, r1, r2, H, C)
    T, U, Wc, X = H['类型'], H['属性号'], H['项号'], H['笔数']
    wi, wig = wy('流_项年属月键'), wy('流_项归年属月键')
    for r in range(r1, r2 + 1):
        t, u, w = f'${T}{r}', f'${U}{r}', f'${Wc}{r}'
        helper(ws, f'{X}{r}', R(f'IF({t}<{TYPE_ITEM},0,COUNTIFS({wi},">="&{k_aypm(w, Y, u, M1)},{wi},"<="&{k_aypm(w, Y, u, M2)}))'))
        for kk, c in enumerate(ccol, 1):
            one = f'SUMIFS({wn},{wig},">="&{k_igypm(w, kk, Y, u, M1)},{wig},"<="&{k_igypm(w, kk, Y, u, M2)})'
            money(ws, f'{c}{r}', R(f'IF({t}=0,"",IF({c}${hr}="","",IF({t}>={TYPE_ITEM},IF(${X}{r}=0,0,ROUND({sign(u)}*{one},2)),'
                                   f'{sub_total(c, r, H, end)})))'))
        one = f'SUMIFS({wn},{wi},">="&{k_aypm(w, Y, u, M1)},{wi},"<="&{k_aypm(w, Y, u, M2)})'
        money(ws, f'{TOT}{r}', R(f'IF({t}=0,"",IF({t}>={TYPE_ITEM},IF(${X}{r}=0,0,ROUND({sign(u)}*{one},2)),{sub_total(TOT, r, H, end)}))'),
              bold=True)
        money(ws, f'{IC_UNL}{r}', R(f'IF({t}=0,"",ROUND({TOT}{r}-SUM({ccol[0]}{r}:{ccol[-1]}{r}),2))'))
    cand_cf(ws, f'A{r1}:{IC_UNL}{r2}', f'${T}{r1}')
    # —— 提示（第 13 行）——
    tr = ITEMCO['提示行']
    sec_sum = {A_IN: SR['经营收入'], A_OUT: SR['经营支出'], A_CAP: SR['投资建设'], A_FLOW: SR['往来净额'], A_TBD: SR['待分类'],
               I9: SR['内部净额']}
    mis = '+'.join(f'(ROUND(SUMIFS(${TOT}${r1}:${TOT}${r2},${T}${r1}:${T}${r2},{TYPE_SEC},${U}${r1}:${U}${r2},{p})-${TOT}${rr},2)<>0)'
                   for p, rr in sec_sum.items())
    direct = f'SUMIFS({wn},{wmm},">="&({Y}*100+{M1}),{wmm},"<="&({Y}*100+{M2}))'
    unl = f'${IC_UNL}${SR["净增加"]}'
    unl_any = 'OR(' + ','.join(f'ROUND(${IC_UNL}${r},2)<>0' for r in SR.values()) + ')'
    tbd = f'${TOT}${SR["待分类"]}'
    tipf = (f'IF(COUNT(分类_候选键)>{N_CLS},"※ 收费项、大类太多：本表只放得下 {N_CLS} 行，后面的没列出来（下面各段会比上面汇总少）",'
            f'IF(({mis})>0,"※ 下面明细跟上面汇总对不上 "&({mis})&" 段：请告诉我们",'
            f'IF(ROUND(${TOT}${SR["净增加"]}-{direct},2)<>0,"※ 资金净增加跟台帐对不上：请告诉我们",'
            f'IF({unl_any},"△ 有的流水归属公司没认出（账户的公司不在基础资料②），资金净增加里有 "&IF({unl}<0,"-","")&{money_txt(unl)}&'
            f'" 不在哪家公司：去③改账户的公司",'
            f'IF(ROUND({tbd},2)<>0,"△ 待分类净额 "&IF({tbd}<0,"-","")&{money_txt(tbd)}&'
            f'"：收费项待定、没填或认不出，去【资金台帐】选收费项（或在基础资料④定属性）",'
            f'"✓ 下面明细合计跟上面汇总一致；合计列跟【收支分类报表】同期全集团一样")))))')
    merge_put(ws, f'A{tr}:{TOT}{tr}', R(tipf), F_TXT, FILL_NONE, align=AL)
    cf_flags(ws, f'A{tr}', f'A{tr}')
    ws.row_dimensions[tr].height = 22
    hide(ws, *H.values(), 'AP', 'AQ')
    hide(ws, *ccol[n_co:])
    ws.freeze_panes = f'{ccol[0]}{r1}'


# ═════════════════════════════ 收支报表 ═════════════════════════════
IE_NO, IE_GATE = 'Q', 'Y'                                  # 块 1～3 每行的收付人号/收费项号；块 2 每行本年笔数（私有）
IE_MWIN = dict(窗首='AA', 窗末='AB')                       # 第 m+1 行＝m 月的窗口（私有）


def build_ie(ws):
    C, RK = IE['列'], IE['排名']
    E = {k: ab(v) for k, v in IE_EFF.items()}
    S = {k: ab(v) for k, v in IE_SEL.items()}
    Y = E['年份']
    mcol = [CL(CI(C['月1']) + m - 1) for m in range(1, 13)]
    FY, PC = C['全年'], C['占比']
    assert CI(FY) == CI(C['月1']) + 12 and CI(PC) == CI(FY) + 1
    b1, b2, b3, b4 = IE['块1'], IE['块2'], IE['块3'], IE['块4']
    widths(ws, {C['序号']: 6, C['项目']: 30, FY: 15, PC: 9})
    for c in mcol:
        ws.column_dimensions[c].width = 12.5
    tip = ('经营口径（跟【收支分类报表】一样）：资金收入＝属性「收入」的行（收入－支出，退款冲减），资金支出＝属性「支出」的行（支出－收入）；'
           '投资建设、往来、待分类、内部转账不在这张表（看【收支分类报表】）。一、收入按收付人：本年收入最多的前 80 个；二、支出按收费项；'
           '三、支出按收付人：前 40 个；「其余」「未列示」「没填或不在清单」是补差，每块合计＝总额。四、按收付人类型（基础资料⑤的「类型」）的全年合计。'
           '名单来自【基础资料】：要单独列出来就去那边加收付人/收费项。年份空着＝管理年度。全部自动算，不用填。')
    title(ws, '=""', PC, tip, h2=58)
    pause_title(ws, f'{group_name()}&" "&{Y}&" 年收支报表（经营口径）"')
    # —— 生效值、月窗口（隐藏）——
    f = dict(年份=yr_eff(S['年份']), 窗首=win_first(f'{Y}*100+1'), 窗末=win_last(f'{Y}*100+12', E['窗首']))
    for k, v in f.items():
        helper(ws, IE_EFF[k], R(v))
        helper(ws, f'{IE_GATE}{IE_EFF[k][1:]}', k)
    W1, W2 = IE_MWIN['窗首'], IE_MWIN['窗末']
    helper(ws, f'{W1}1', '月窗首')
    helper(ws, f'{W2}1', '月窗末')
    for m in range(1, 13):
        helper(ws, f'{W1}{m + 1}', R(win_first(f'{Y}*100+{m}')))
        helper(ws, f'{W2}{m + 1}', R(win_last(f'{Y}*100+{m}', f'{W1}{m + 1}')))
    wm = lambda nm, m: win(nm, f'${W1}${m + 1}', f'${W2}${m + 1}')
    wy = lambda nm: win(nm, E['窗首'], E['窗末'])
    # —— 第 3、4 行 ——
    selector(ws, 'A3', '年份', IE_SEL['年份'], kind='year', prompt='填 2026 或 26；空着＝管理年度')
    merge_put(ws, 'C3:E3', by_year(S['年份'], Y), F_NOTE, FILL_NONE, align=ASHRINK)
    merge_put(ws, 'F3:G3', '本年经营净额', F_KPI_L, FILL_KPI, align=ACW)
    merge_put(ws, 'H3:I3', R(f'{FY}{IE["盈亏行"]}'), F_KPI_V, FILL_KPI, fmt=ACC)
    cf_flags(ws, 'C3', 'C3')
    home_link(ws, FY + '3')
    ws.row_dimensions[3].height = 24
    hdr = [(C['序号'], '序号'), (C['项目'], '项目')] + [(c, f'{m}月') for m, c in enumerate(mcol, 1)] + [(FY, '全年'), (PC, '占比')]
    header(ws, IE['表头行'], hdr, '报表', height=24)
    # —— 收支盈亏 ——
    r5 = IE['盈亏行']
    txt(ws, f'{C["序号"]}{r5}', None, fill_=FILL_KPI)
    txt(ws, f'{C["项目"]}{r5}', '收支盈亏（经营收入－经营支出）', bold=True, fill_=FILL_KPI)
    for c in mcol:
        money(ws, f'{c}{r5}', R(f'ROUND({c}{b1["总额"]}-{c}{b2["总额"]},2)'), bold=True, fill_=FILL_KPI)
    money(ws, f'{FY}{r5}', R(f'ROUND(SUM({mcol[0]}{r5}:{mcol[-1]}{r5}),2)'), fill_=FILL_KPI, font=F_RED)
    put(ws, f'{PC}{r5}', None, F_AUTO, FILL_KPI)
    cf_neg_red(ws, f'{mcol[0]}{r5}:{FY}{r5}', f'{mcol[0]}{r5}')

    def total_row(r, no, label, p, sg):
        txt(ws, f'{C["序号"]}{r}', no, bold=True, fill_=FILL_SECT, align=AC)
        txt(ws, f'{C["项目"]}{r}', label, bold=True, fill_=FILL_SECT)
        for m, c in enumerate(mcol, 1):
            v = f'SUMIFS({wm("流_净", m)},{wm("流_月码", m)},{Y}*100+{m},{wm("流_属性号", m)},{p})'
            money(ws, f'{c}{r}', R(f'ROUND({"-" if sg < 0 else ""}{v},2)'), bold=True, fill_=FILL_SECT)
        money(ws, f'{FY}{r}', R(f'ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2)'), bold=True, fill_=FILL_SECT)
        pct(ws, f'{PC}{r}', R(f'IF(ROUND({FY}{r},2)=0,"",1)'), fill_=FILL_SECT)

    def share_cell(r, tot_r):
        pct(ws, f'{PC}{r}', R(f'IF(OR(N({FY}{r})=0,ROUND(${FY}${tot_r},2)=0),"",{FY}{r}/${FY}${tot_r})'))

    def rest_row(r, tot_r, first, last, unk_r, label):
        """其余＝总额－明细－没填（每月、全年都补差）"""
        txt(ws, f'{C["项目"]}{r}', label, font=F_NOTE)
        put(ws, f'{C["序号"]}{r}', None, F_AUTO, FILL_NONE)
        for c in mcol + [FY]:
            unk = f'-{c}{unk_r}' if unk_r else ''
            money(ws, f'{c}{r}', R(f'ROUND({c}{tot_r}-SUM({c}{first}:{c}{last}){unk},2)'))
        share_cell(r, tot_r)

    def unk_row(r, tot_r, p, sg, label):
        txt(ws, f'{C["项目"]}{r}', label, font=F_NOTE)
        put(ws, f'{C["序号"]}{r}', None, F_AUTO, FILL_NONE)
        for m, c in enumerate(mcol, 1):
            v = f'SUMIF({wm("流_人年属月键", m)},{k_aypm(0, Y, p, m)},{wm("流_净", m)})'
            money(ws, f'{c}{r}', R(f'ROUND({"-" if sg < 0 else ""}{v},2)'))
        money(ws, f'{FY}{r}', R(f'ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2)'))
        share_cell(r, tot_r)

    def party_rows(blk, p, sg, kcol, n):
        first, last = blk['首行'], blk['末行']
        assert last - first + 1 == n
        for r in range(first, last + 1):
            q = f'${IE_NO}{r}'
            helper(ws, f'{IE_NO}{r}', R(f'N(INDEX(${kcol}$2:${kcol}${n + 1},ROWS(${IE_NO}${first}:{IE_NO}{r})))'))
            put(ws, f'{C["序号"]}{r}', R(f'IF({q}=0,"",ROWS(${IE_NO}${first}:{IE_NO}{r}))'), F_AUTO, FILL_NONE, fmt=INT, align=AC)
            txt(ws, f'{C["项目"]}{r}', R(f'IF({q}=0,"","　"&INDEX(人_名称,{q}))'), align=ASHRINK)
            for m, c in enumerate(mcol, 1):
                v = f'SUMIF({wm("流_人年属月键", m)},{k_aypm(q, Y, p, m)},{wm("流_净", m)})'
                money(ws, f'{c}{r}', R(f'IF({q}=0,"",ROUND({"-" if sg < 0 else ""}{v},2))'))
            money(ws, f'{FY}{r}', R(f'IF({q}=0,"",ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2))'), bold=True)
            share_cell(r, blk['总额'])
    # —— 块 1：资金收入（按收付人）——
    total_row(b1['总额'], '一', '资金收入合计（经营收入）· 按收付人', A_IN, 1)
    n_in = f'COUNTIF(${RK["年收入"]}$2:${RK["年收入"]}${N_PARTY + 1},"<>0")'
    listed1 = f'COUNTIF(${IE_NO}${b1["首行"]}:${IE_NO}${b1["末行"]},">0")'
    rest_row(b1['其余'], b1['总额'], b1['首行'], b1['末行'], b1['没填'], None)
    ws[f'{C["项目"]}{b1["其余"]}'] = R(f'"　其余 "&MAX(0,{n_in}-{listed1})&" 个收付人（前 {IE_TOP_IN} 名以外的、全年收入不是正数的）"')
    unk_row(b1['没填'], b1['总额'], A_IN, 1, '　收付人没填或不在清单')
    party_rows(b1, A_IN, 1, RK['收第k'], IE_TOP_IN)
    # —— 块 2：资金支出（按收费项）——
    total_row(b2['总额'], '二', '资金支出合计（经营支出）· 按收费项', A_OUT, -1)
    rest_row(b2['未列示'], b2['总额'], b2['首行'], b2['末行'], None, f'　未列示（收费项超过 {IE_ITEMS} 个、认不出的）')
    wi = wy('流_项年属月键')
    for r in range(b2['首行'], b2['末行'] + 1):
        q, g = f'${IE_NO}{r}', f'${IE_GATE}{r}'
        k = f'ROWS(${IE_NO}${b2["首行"]}:{IE_NO}{r})'
        helper(ws, f'{IE_NO}{r}', R(f'IF({k}>项_序数{A_OUT},0,MOD(SMALL(项_序键{A_OUT},{k}),1000))'))
        helper(ws, f'{IE_GATE}{r}', R(f'IF({q}=0,0,COUNTIFS({wi},">="&{k_aypm(q, Y, A_OUT, 1)},{wi},"<="&{k_aypm(q, Y, A_OUT, 12)}))'))
        put(ws, f'{C["序号"]}{r}', R(f'IF({q}=0,"",{k})'), F_AUTO, FILL_NONE, fmt=INT, align=AC)
        txt(ws, f'{C["项目"]}{r}', R(f'IF({q}=0,"","　"&INDEX(项_名称,{q}))'), align=ASHRINK)
        for m, c in enumerate(mcol, 1):
            v = f'SUMIF({wm("流_项年属月键", m)},{k_aypm(q, Y, A_OUT, m)},{wm("流_净", m)})'
            money(ws, f'{c}{r}', R(f'IF({q}=0,"",IF({g}=0,0,ROUND(-{v},2)))'))
        money(ws, f'{FY}{r}', R(f'IF({q}=0,"",ROUND(SUM({mcol[0]}{r}:{mcol[-1]}{r}),2))'), bold=True)
        share_cell(r, b2['总额'])
    # —— 块 3：资金支出（按收付人）——
    total_row(b3['总额'], '三', '资金支出合计（经营支出）· 按收付人', A_OUT, -1)
    n_out = f'COUNTIF(${RK["年支出"]}$2:${RK["年支出"]}${N_PARTY + 1},"<>0")'
    listed3 = f'COUNTIF(${IE_NO}${b3["首行"]}:${IE_NO}${b3["末行"]},">0")'
    rest_row(b3['其余'], b3['总额'], b3['首行'], b3['末行'], b3['没填'], None)
    ws[f'{C["项目"]}{b3["其余"]}'] = R(f'"　其余 "&MAX(0,{n_out}-{listed3})&" 个收付人（前 {IE_TOP_OUT} 名以外的、全年支出不是正数的）"')
    unk_row(b3['没填'], b3['总额'], A_OUT, -1, '　收付人没填或不在清单')
    party_rows(b3, A_OUT, -1, RK['支第k'], IE_TOP_OUT)
    for blk in (b1, b2, b3):
        cf_neg_red(ws, f'{mcol[0]}{blk["总额"]}:{FY}{blk["末行"]}', f'{mcol[0]}{blk["总额"]}')
    # —— 排名表（隐藏，第 2..501 行＝收付人 1..500）——
    for k, t in (('人号', '人号'), ('年收入', '年收入'), ('收键', '收键'), ('年支出', '年支出'), ('支键', '支键'), ('收第k', '收第k'),
                 ('支第k', '支第k')):
        helper(ws, f'{RK[k]}1', t)
    wn, wk = wy('流_净'), wy('流_人年属月键')
    for i in range(1, N_PARTY + 1):
        r = i + 1
        rr = f'${RK["人号"]}{r}'
        helper(ws, f'{RK["人号"]}{r}', i)
        yv = lambda p: f'SUMIFS({wn},{wk},">="&{k_aypm(rr, Y, p, 1)},{wk},"<="&{k_aypm(rr, Y, p, 12)})'
        helper(ws, f'{RK["年收入"]}{r}', R(f'IF(INDEX(人_有效,{rr})<>1,0,ROUND({yv(A_IN)},2))'))
        helper(ws, f'{RK["收键"]}{r}', R(f'IF({RK["年收入"]}{r}>0,{RK["年收入"]}{r}+{rr}/1000000,"")'))
        helper(ws, f'{RK["年支出"]}{r}', R(f'IF(INDEX(人_有效,{rr})<>1,0,ROUND(-{yv(A_OUT)},2))'))
        helper(ws, f'{RK["支键"]}{r}', R(f'IF({RK["年支出"]}{r}>0,{RK["年支出"]}{r}+{rr}/1000000,"")'))
        for kcol, key, n in ((RK['收第k'], RK['收键'], IE_TOP_IN), (RK['支第k'], RK['支键'], IE_TOP_OUT)):
            if i <= n:
                rg = f'${key}$2:${key}${N_PARTY + 1}'
                helper(ws, f'{kcol}{r}', R(f'IFERROR(MATCH(LARGE({rg},{rr}),{rg},0),0)'))
    # —— 块 4：按收付人类型（全年）——
    section(ws, b4['标题'], C['序号'], 'G', '四、按收付人类型（全年；类型在基础资料⑤）')
    h4 = dict(A='序号', B='收付人类型', C='经营收入', D='占收入', E='经营支出', F='占支出', G='收支盈亏')
    header(ws, b4['表头'], list(h4.items()), '报表', height=24)
    sin, sout = f'${RK["年收入"]}$2:${RK["年收入"]}${N_PARTY + 1}', f'${RK["年支出"]}$2:${RK["年支出"]}${N_PARTY + 1}'
    tr4 = b4['合计']
    for j, t in enumerate(PARTY_TYPES):
        r = b4['首行'] + j
        put(ws, f'A{r}', j + 1, F_AUTO, FILL_NONE, align=AC)
        txt(ws, f'B{r}', t)
        money(ws, f'C{r}', R(f'ROUND(SUMIF(人_类型,$B{r},{sin}),2)'))
        pct(ws, f'D{r}', R(f'IF(ROUND($C${tr4},2)=0,"",C{r}/$C${tr4})'))
        money(ws, f'E{r}', R(f'ROUND(SUMIF(人_类型,$B{r},{sout}),2)'))
        pct(ws, f'F{r}', R(f'IF(ROUND($E${tr4},2)=0,"",E{r}/$E${tr4})'))
        money(ws, f'G{r}', R(f'ROUND(C{r}-E{r},2)'), bold=True)
    assert b4['首行'] + len(PARTY_TYPES) == tr4
    merge_put(ws, f'A{tr4}:B{tr4}', '合计（＝上面一、二的总额）', F_RED, FILL_KPI)
    money(ws, f'C{tr4}', R(f'{FY}{b1["总额"]}'), fill_=FILL_KPI, font=F_RED)
    pct(ws, f'D{tr4}', R(f'IF(ROUND(C{tr4},2)=0,"",1)'), fill_=FILL_KPI)
    money(ws, f'E{tr4}', R(f'{FY}{b2["总额"]}'), fill_=FILL_KPI, font=F_RED)
    pct(ws, f'F{tr4}', R(f'IF(ROUND(E{tr4},2)=0,"",1)'), fill_=FILL_KPI)
    money(ws, f'G{tr4}', R(f'ROUND(C{tr4}-E{tr4},2)'), fill_=FILL_KPI, font=F_RED)
    ru = tr4 + 1                                             # 私有：没归到 9 类的
    merge_put(ws, f'A{ru}:B{ru}', '其中没归到上面几类的（收付人没填/不在清单、没写类型）', F_NOTE, FILL_NONE, align=ASHRINK)
    for c in 'CEG':
        money(ws, f'{c}{ru}', R(f'ROUND({c}{tr4}-SUM({c}{b4["首行"]}:{c}{tr4 - 1}),2)'))
    for c, d in (('D', 'C'), ('F', 'E')):
        pct(ws, f'{c}{ru}', R(f'IF(OR(N({d}{ru})=0,ROUND({d}{tr4},2)=0),"",{d}{ru}/{d}{tr4})'))
    hide(ws, IE_NO, IE_GATE, *RK.values(), 'Z', W1, W2)
    ws.freeze_panes = f'{mcol[0]}{IE["盈亏行"]}'


# ═════════════════════════════ 任意时间段收支报表 ═════════════════════════════
SPAN_PRIV = dict(号='Z', 笔数='AA')                         # 各账户第 8..47 行（私有）
SPAN_NO = dict(收入按收费项='AB', 收入按收付人='AC', 支出按收费项='AD', 支出按收付人='AE')   # 四块每行的号（私有）


def build_span(ws):
    E = {k: ab(v) for k, v in SPAN_EFF.items()}
    A_ = SPAN_ACC['列']
    K = SPAN_KPI
    D1, D2 = 'P_期起', 'P_期止'
    widths(ws, dict(A=26, B=15, C=15, D=15, E=18, F=15, G=15, H=15, I=18, J=15, K=13, L=15, M=18, N=15, O=9))
    tip = ('任意一段日期（可以跨月、跨年，最多到日期上限）：上面是这段时间的经营收入、支出等（经营口径，跟【收支分类报表】一样）和各账户的'
           '期初、收入、支出、内部转入/转出、期末（账户口径：收入、支出不含内部转账；期末＝期初＋收入－支出＋内部转入－内部转出；余额按日期算）；'
           '下面四块：收入按收费项、收入按收付人（前 80 名）、支出按收费项、支出按收付人（前 40 名），「未列示」是补差。'
           '黄格：开始空着＝最后一笔流水那年的 1 月 1 日，结束空着＝最后一笔流水的日期。全部自动算，不用填。')
    title(ws, '=""', 'O', tip, h2=58)
    pause_title(ws, f'{group_name()}&" 任意时间段收支报表（"&{d_txt(D1)}&"～"&IF({D2}<{D1},"※ 结束早于开始",{d_txt(D2)})&"）"')
    # —— 窗口（隐藏）——
    f = dict(窗首=win_first(ym_of(D1)), 窗末=win_last(ym_of(D2), E['窗首']), 前窗末=win_last(ym_of(f'({D1}-1)')))
    for k, v in f.items():
        helper(ws, SPAN_EFF[k], R(v))
        helper(ws, f'X{SPAN_EFF[k][1:]}', k)
    w = lambda nm: win(nm, E['窗首'], E['窗末'])
    w0 = lambda nm: win(nm, 1, E['前窗末'])
    # —— 第 3、4 行：选择格＋大数 ——
    selector(ws, 'A3', '开始日期', SPAN_SEL['开始'], kind='date', prompt='空着＝最后一笔流水那年的 1 月 1 日')
    selector(ws, 'A4', '结束日期', SPAN_SEL['结束'], kind='date', prompt='空着＝最后一笔流水的日期')
    sp = lambda p: f'SUMIFS({w("流_净")},{w("流_日期")},">="&{D1},{w("流_日期")},"<="&{D2},{w("流_属性号")},{p})'
    kv = dict(经营收入=sp(A_IN), 经营支出=f'-{sp(A_OUT)}', 经营净额=f'{K["经营收入"]}-{K["经营支出"]}', 投资建设=f'-{sp(A_CAP)}',
              往来净额=sp(A_FLOW), 待分类=sp(A_TBD), 内部净额=sp(I9))
    lab = dict(经营收入='经营收入', 经营支出='经营支出', 经营净额='经营净额', 投资建设='投资建设', 往来净额='往来净额', 待分类='待分类净额',
               内部净额='内部转账净额', 天数笔数='期间')
    for k, cell in K.items():
        lc = CL(CI(re.match(r'[A-Z]+', cell).group()) - 1) + cell[len(re.match(r'[A-Z]+', cell).group()):]
        put(ws, lc, lab[k], F_KPI_L, FILL_KPI, align=ACW)
        if k == '天数笔数':
            cnt = f'COUNTIFS({w("流_日期")},">="&{D1},{w("流_日期")},"<="&{D2},{w("流_有效")},1)'
            put(ws, cell, R(f'IF({D2}<{D1},"※ 结束早于开始，请改日期","共 "&({D2}-{D1}+1)&" 天 "&{cnt}&" 笔")'), F_KPI_L, FILL_KPI,
                align=ACSHRINK)
        else:
            put(ws, cell, R(f'ROUND({kv[k]},2)'), F_KPI_V, FILL_KPI, fmt=ACC, align=AC)
    cf_flags(ws, K['天数笔数'], K['天数笔数'])
    tr = SPAN_ACC['合计行']
    put(ws, 'K3', '资金净增加', F_KPI_L, FILL_KPI, align=ACW)
    put(ws, 'L3', R(f'ROUND({K["经营净额"]}-{K["投资建设"]}+{K["往来净额"]}+{K["待分类"]}+{K["内部净额"]},2)'), F_KPI_V, FILL_KPI,
        fmt=ACC, align=AC)
    put(ws, 'K4', '核对（应为 0）', F_KPI_L, FILL_KPI, align=ACW)
    put(ws, 'L4', R(f'ROUND(L3-({A_["期末余额"]}{tr}-{A_["期初余额"]}{tr}),2)'), F_KPI_V, FILL_KPI, fmt=ACC, align=AC)
    ws.conditional_formatting.add('L4', FormulaRule(formula=['AND(ISNUMBER(L4),ROUND(L4,2)<>0)'],
                                                    font=Font(name=YH, bold=True, color='FFC00000'), fill=cf_bg('FFFFD7D7')))
    for k in ('经营净额', '待分类'):
        cf_neg_red(ws, K[k], K[k])
    home_link(ws, 'N3')
    ws.row_dimensions[3].height = ws.row_dimensions[4].height = 24
    s1, s2 = ab(SPAN_SEL['开始']), ab(SPAN_SEL['结束'])
    note = (f'"按 "&{d_txt(D1)}&"～"&{d_txt(D2)}&" 算"&IF({raw(s1)},"","（开始空着＝"&YEAR(P_最后日期)&" 年 1 月 1 日）")&'
            f'IF({raw(s2)},"","（结束空着＝最后一笔流水）")')
    merge_put(ws, 'A5:O5', R(f'IF({dsel_bad(s1)},"※ 开始日期看不懂，按 "&{d_txt(D1)}&" 算；",IF({dsel_bad(s2)},"※ 结束日期看不懂，按 "&'
                             f'{d_txt(D2)}&" 算；",""))&{note}'), F_NOTE, FILL_NONE, align=AL)
    cf_flags(ws, 'A5', 'A5')
    # —— 各账户（第 6～48 行）——
    section(ws, SPAN_ACC['标题行'], 'A', 'G', '各账户（账户口径：收入、支出不含内部转账；期末＝期初＋收入－支出＋内部转入－内部转出）')
    hdr = dict(账户='账户', 期初余额='期初余额\n（开始那天之前）', 收入='收入\n（不含内部）', 支出='支出\n（不含内部）', 内部转入='内部转入',
               内部转出='内部转出', 期末余额='期末余额\n（结束那天）')
    header(ws, SPAN_ACC['表头行'], [(A_[k], hdr[k]) for k in A_], '报表', height=34)
    P = SPAN_PRIV
    helper(ws, f'{P["号"]}{SPAN_ACC["表头行"]}', '账户号')
    helper(ws, f'{P["笔数"]}{SPAN_ACC["表头行"]}', '期间笔数')
    r1 = SPAN_ACC['首行']
    wk, wk0 = w('流_账日键'), w0('流_账日键')
    for r in range(r1, tr):
        k = f'ROWS(${P["号"]}${r1}:{P["号"]}{r})'
        a, n = f'${P["号"]}{r}', f'${P["笔数"]}{r}'
        helper(ws, f'{P["号"]}{r}', R(f'IF(INDEX(账户_有效,{k})=1,{k},0)'))
        helper(ws, f'{P["笔数"]}{r}', R(f'IF(OR({a}=0,{D2}<{D1}),0,COUNTIFS({wk},">="&({a}*{K_DAY}+{D1}),{wk},"<="&({a}*{K_DAY}+{D2})))'))
        txt(ws, f'{A_["账户"]}{r}', R(f'IF({a}=0,"",INDEX(账户_名称,{a}))'), align=ASHRINK)
        money(ws, f'{A_["期初余额"]}{r}', R(f'IF({a}=0,"",ROUND(INDEX(账户_期初,{a})+SUMIFS({w0("流_净")},{wk0},">="&{a}*{K_DAY},'
                                           f'{wk0},"<"&({a}*{K_DAY}+{D1})),2))'))
        for kk, val, cond in (('收入', '流_收', f'"<>{I9}"'), ('支出', '流_支', f'"<>{I9}"'), ('内部转入', '流_收', I9),
                              ('内部转出', '流_支', I9)):
            money(ws, f'{A_[kk]}{r}', R(f'IF({a}=0,"",IF({n}=0,0,ROUND(SUMIFS({w(val)},{wk},">="&({a}*{K_DAY}+{D1}),{wk},"<="&({a}*{K_DAY}+{D2}),'
                                       f'{w("流_属性号")},{cond}),2)))'))
        c = lambda x: f'{A_[x]}{r}'
        money(ws, c('期末余额'), R(f'IF({a}=0,"",ROUND({c("期初余额")}+{c("收入")}-{c("支出")}+{c("内部转入")}-{c("内部转出")},2))'),
              bold=True)
    put(ws, f'{A_["账户"]}{tr}', '合计', F_RED, FILL_KPI, align=AC)
    for kk in ('期初余额', '收入', '支出', '内部转入', '内部转出', '期末余额'):
        put(ws, f'{A_[kk]}{tr}', R(f'ROUND(SUM({A_[kk]}{r1}:{A_[kk]}{tr - 1}),2)'), F_RED, FILL_KPI, fmt=ACC, align=AR)
    for kk in ('期初余额', '期末余额'):
        cf_neg_red(ws, f'{A_[kk]}{r1}:{A_[kk]}{tr}', f'{A_[kk]}{r1}')
    # —— 排名表（隐藏，第 2..501 行＝收付人 1..500；按 流_期人键）——
    RK = SPAN_RANK
    for k in RK:
        helper(ws, f'{RK[k]}1', k)
    wpk, wn = w('流_期人键'), w('流_净')
    for i in range(1, N_PARTY + 1):
        r = i + 1
        rr = f'${RK["人号"]}{r}'
        helper(ws, f'{RK["人号"]}{r}', i)
        helper(ws, f'{RK["期收入"]}{r}', R(f'IF(INDEX(人_有效,{rr})<>1,0,ROUND(SUMIF({wpk},{rr}*10+{A_IN},{wn}),2))'))
        helper(ws, f'{RK["收键"]}{r}', R(f'IF({RK["期收入"]}{r}>0,{RK["期收入"]}{r}+{rr}/1000000,"")'))
        helper(ws, f'{RK["期支出"]}{r}', R(f'IF(INDEX(人_有效,{rr})<>1,0,ROUND(-SUMIF({wpk},{rr}*10+{A_OUT},{wn}),2))'))
        helper(ws, f'{RK["支键"]}{r}', R(f'IF({RK["期支出"]}{r}>0,{RK["期支出"]}{r}+{rr}/1000000,"")'))
        for kcol, key, n in ((RK['收第k'], RK['收键'], SPAN_BLK['行数']['收入按收付人']), (RK['支第k'], RK['支键'], SPAN_BLK['行数']['支出按收付人'])):
            if i <= n:
                rg = f'${key}$2:${key}${N_PARTY + 1}'
                helper(ws, f'{kcol}{r}', R(f'IFERROR(MATCH(LARGE({rg},{rr}),{rg},0),0)'))
    # —— 四块并排 ——
    B = SPAN_BLK
    t0, h0, tot, unl, f0 = B['标题行'], B['表头行'], B['合计行'], B['未列示行'], B['首行']
    spec = dict(收入按收费项=('收入 · 按收费项（经营收入）', '收费项', A_IN, 1, K['经营收入'], '项'),
                收入按收付人=(f'收入 · 按收付人（前 {SPAN_BLK["行数"]["收入按收付人"]} 名）', '收付人', A_IN, 1, K['经营收入'], '人'),
                支出按收费项=('支出 · 按收费项（经营支出）', '收费项', A_OUT, -1, K['经营支出'], '项'),
                支出按收付人=(f'支出 · 按收付人（前 {SPAN_BLK["行数"]["支出按收付人"]} 名）', '收付人', A_OUT, -1, K['经营支出'], '人'))
    wpi, wpp, wpd = w('流_项号'), w('流_属性号'), w('流_日期')
    for blk, (cn, ca, cs) in B['块'].items():
        ttl, hname, p, sg, kcell, kind = spec[blk]
        nrow = B['行数'][blk]
        hc = SPAN_NO[blk]
        section(ws, t0, cn, cs, ttl)
        header(ws, h0, [(cn, hname), (ca, '金额'), (cs, '占比')], '报表', height=22)
        helper(ws, f'{hc}{h0}', blk + '·号')
        put(ws, f'{cn}{tot}', '合计', F_RED, FILL_KPI, align=AC)
        put(ws, f'{ca}{tot}', R(f'{ab(kcell)}'), F_RED, FILL_KPI, fmt=ACC, align=AR)
        pct(ws, f'{cs}{tot}', R(f'IF(ROUND({ca}{tot},2)=0,"",1)'), fill_=FILL_KPI)
        last = f0 + nrow - 1
        txt(ws, f'{cn}{unl}', '未列示的收费项' if kind == '项' else '其余收付人、没填或不在清单的', font=F_NOTE, align=ASHRINK)
        money(ws, f'{ca}{unl}', R(f'ROUND({ca}{tot}-SUM({ca}{f0}:{ca}{last}),2)'))
        pct(ws, f'{cs}{unl}', R(f'IF(OR(N({ca}{unl})=0,ROUND(${ca}${tot},2)=0),"",{ca}{unl}/${ca}${tot})'))
        for r in range(f0, last + 1):
            k = f'ROWS(${hc}${f0}:{hc}{r})'
            q = f'${hc}{r}'
            if kind == '项':
                helper(ws, f'{hc}{r}', R(f'IF({k}>项_序数{p},0,MOD(SMALL(项_序键{p},{k}),1000))'))
                txt(ws, f'{cn}{r}', R(f'IF({q}=0,"",INDEX(项_名称,{q}))'), align=ASHRINK)
                v = f'SUMIFS({wn},{wpi},{q},{wpp},{p},{wpd},">="&{D1},{wpd},"<="&{D2})'
                money(ws, f'{ca}{r}', R(f'IF({q}=0,"",ROUND({"-" if sg < 0 else ""}{v},2))'))
            else:
                src = RK['收第k'] if p == A_IN else RK['支第k']
                val = RK['期收入'] if p == A_IN else RK['期支出']
                helper(ws, f'{hc}{r}', R(f'N(INDEX(${src}$2:${src}${nrow + 1},{k}))'))
                txt(ws, f'{cn}{r}', R(f'IF({q}=0,"",INDEX(人_名称,{q}))'), align=ASHRINK)
                money(ws, f'{ca}{r}', R(f'IF({q}=0,"",INDEX(${val}$2:${val}${N_PARTY + 1},{q}))'))
            pct(ws, f'{cs}{r}', R(f'IF(OR({q}=0,ROUND(${ca}${tot},2)=0),"",{ca}{r}/${ca}${tot})'))
        cf_neg_red(ws, f'{ca}{tot}:{ca}{last}', f'{ca}{tot}')
    hide(ws, *RK.values(), 'X', 'Y', *SPAN_PRIV.values(), *SPAN_NO.values())
    ws.freeze_panes = 'B8'


def build(wb, ctx):
    build_cls(wb[SH_CLS])
    build_itemco(wb[SH_ITEMCO], len(ctx['companies']))
    build_ie(wb[SH_IE])
    build_span(wb[SH_SPAN])
