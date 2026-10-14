# -*- coding: utf-8 -*-
"""A074 集团多账户资金台帐 · 【核对表】【数据整理记录】【使用说明】（设计.md §7；v3.1 提速改动见 §15.1.5、§15.2、§15.3.1、§15.5）。

【核对表】（layout.CHK：4 行表头，5 行起；列 序号｜类别｜检查项｜结果｜判断｜说明；隐藏计数格 CHK_CNT 在 H 列，标签在 I 列）
  判断三级：✓（通过/没有）、△（看一眼、最好补）、✗（要改 / 算不平）；「— …」是不算的（录入模式中、备用子表没选账户）。第 3 行汇总。
  一、勾稽：每张报表交出来的格子（layout 的 日报_/月报_/公汇_/分类_/分项_/收支_/时段_/年末_/首页_ 名称格）拿报表自己选的年份、日期，
     跟 _流 直接按口径算的比（不开窗口，整列 SUMIFS：报表开了窗口，这里不开，窗口算漏了也查得出）；各项差放隐藏帮助格（K 列起），
     结果＝Σ|差|，说明里点出第一个差的是哪一项。另有：恒等式、各报表互相对得上、内部转账本年净额＝0、规则推断自检。
  二、子表（26 张＋备用 2＋账户查询）：每张一行——子表的余额 J4、月份表合计行 N19～R19、笔数 K4、明细收－支（自己 SUM 本表 G、H）
     跟 _流 直接算的比（v3.1：只扫那一年／那个月的窗口 layout.win_first/win_last，余额扫 [1, 年底窗口末]）；
     笔数＞容量 ✗；26 张子表认不出账户（Z1=0）△；子表被删 ✗。引用子表的格一律包 IFERROR。
  三、待办：※/△ 行数、认不出的名字、按摘要认的、待分类、没配上、收付人、日期顺序（约）、重复、跳期、暂停户还在记（约）、
     负余额账户、基础资料核对列。v3.1 计数口径（设计 §15.3.1）：跳期、日期顺序数 流_顺序，重复数 流_提本（准数），
     按摘要认＝COUNTIFS(流_项推,">0",流_对方号,0,流_有效,1)，认不出的写法从 流_问本 取；「约」＝被排在前面的提示盖住的数不到。
  四、容量和结构：台帐已用／能录（P_可录）、次年流水、旧年流水没删、**台帐分段**（设计 §15.1.5：锚列/表头被删、格子被挪乱、
     有行没算进来（列出前 5 处在台帐第几行）、整段删掉、借位）、台帐列数、显示列缺公式或被盖掉（△，只当缺公式用，错位归「台帐分段」）、
     基础资料中间插的行没带公式（复核 compat-3）、基础资料各块容量外有字、人规则超 100、规Q 超 150、规则没起作用、大类超 40、
     收支分类行位、名单用量。
  下面清单（链式 MATCH，隐藏 K 列＝在 _流 第几个位置；台帐行号用 layout.z_row，原文用 layout.lref_z）：
     五 认不出的写法（前 10 种，从前 100 笔这种 ※ 里去重）、六 内部转账没配上、七 ※ 的行、八 按摘要认的、九 可能重复粘贴（各前 CHK_LIST 笔）。
  计数格（CHK_CNT，首页提醒条、台帐汇总条直接引用）：
     问题 H1＝COUNTIF(流_问题,"※*")＋P_没算＋MAX(0,−P_笔数差)＋(P_错位>0)＋P_锚坏（设计 §15.1.5；录入模式照算）；
     提醒 H2＝COUNTIF(流_问题,"△*")＋COUNTIFS(流_问题,"",流_配对,1)（录入模式照算；配对停算＝""，没配上的不数）；
     其余 H3～H11 录入模式出 0（pause0）。
  录入模式（设计 §15.2.2，P_录入模式=1）：本表在 PAUSE_SHEETS 里，第 2 行起每个公式以 pause()/pause0() 打头；
     不停的格登记 ws._pause_exempt：H1、H2、第 3 行汇总 C3/F3、「判断」列（显示「— 录入模式中」，只引用本行结果格和帮助格）、
     清单标题（只引用本表计数帮助格）。第 1 行标题换成 layout.PAUSE_TITLE（红字）。
  COUNTIFS 的第一个条件区域一律放 _流 的列（LibreOffice 的 COUNTIFS 会把第一个区域首尾的空格子裁掉）。
  防环：本表只引用名称（流_/段_/块_/P_/台帐_/基_/台k_/报表名称格）、报表/子表/首页交出来的格子；没有别的表引用本表除 CHK_CNT 以外的格。
     段表叶子（段_错位、段_有字）不引用（make 查）：要定位哪一段错位，本表自己在隐藏列按 台k_字段 跟 段_区首/区末 比（只在 P_错位>0 时算）。
【数据整理记录】静态（全是值；v3.1 设计 §15.5）：build_log(ws, ctx, detail=False) 主本只留要点——一 整理口径、二 待确认（带清单的逐笔列在
  那一条下面）、三 账户对照、四 模板这边的调整，最后一行写「五、修改明细 N 条在同一文件夹的《A074_数据整理记录.xlsx》」；
  detail=True（build.py 写单独一本）五部分全有。一张能筛选的表（第 4 行表头、冻结、自动筛选），「部分」一列分段。
【使用说明】静态，自己的话（不抄参考模板）。笔数、条数从 ctx 取。
"""
import datetime as dt
import os
import re
from openpyxl.styles import Font, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from layout import (
    SH_HOME, SH_LEDGER, SH_BASE, SH_DAY, SH_MON, SH_CO, SH_CLS, SH_ITEMCO, SH_IE, SH_SPAN, SH_YE, SH_QRY, SH_CHK, SH_LOG,
    SH_HELP, SUB_SHEETS, SPARE_SHEETS, SUB_CAP, N_QUERY, N_LEDGER, N_BLK, N_ACC, N_CO, N_ITEM, N_PARTY, N_RULE, N_RULE_P,
    N_RULE_Q, N_GROUP, N_CLS, N_SEG, N_FLOW, SEG_ROWS, SEG_SLOTS, SEG_GROW_MAX, CHK, CHK_CNT, CHK_LIST, ATTR_NO, LEDGER_KPI,
    LEDGER_COLS, LEDGER_SHOW, LEDGER_LAST, LEDGER_ANCHOR, BASE_BLOCKS, BASE_COLS, BASE_KEYCOL, BASE_LAST, BASE_CAP, CO1, CLS,
    ITEMCO, SUB_R0, SUB_PARAM, SUB_KPI, SUB_MON, SUB_MON_TOT, SUB_COLS, K_YM, PAIR_DAYS, DUP_WIN, JUMP_DAYS, BLK,
    ITEM_OPEN, R0, HDR, PAUSE_TITLE, q, bref, lrange, esc, next_pos, k_aypm, z_row, lref_z, seg_name, win_first, win_last, win,
    pause, pause0)
from common import (
    YH, F_TXT, F_TXTB, F_AUTO, F_AUTOB, F_HELP, F_NOTE, F_KPI_L, F_SEC, FILL_KPI, FILL_NONE, BD, AC, AL, ALW, AR, ACC, INT, PCT,
    DATE, put, title, header, widths, hide, link, home_link, cf_bg, fill, print_setup)

I1, I2, I3, I4, I5, I9 = (ATTR_NO[k] for k in ('收入', '支出', '投资建设', '往来', '待分类', '内部转账'))
C = CHK['列']
ENTRY = 'P_录入模式=1'
J_ENTRY = '— 录入模式中'
CNT = {k: f'${v[0]}${v[1:]}' for k, v in CHK_CNT.items()}         # 'H1' → '$H$1'
HELP0 = CI('K')                                                  # 各行隐藏帮助格从 K 列起（子表那几行用到 AB）
UNK_COL = dict(位置='AF', 键='AG', 首现='AH', 累计='AI')             # 「认不出的写法」去重用的 100 笔（第 5..104 行）
UNK_N = 100
NEG_COL = dict(余额='AK', 负='AL')                                        # 每个账户的当前余额（第 5..44 行，账户 a 在第 4+a 行）
SEG_COL = dict(错位='AN', 没算='AO', 位置='AP', 段='AQ', 容量外='AR')      # 台帐分段：每段（第 4+k 行）错位累计、没算累计；AP5..AP9 前 5 处、AQ 那几处的段号、AR5 容量外行数
HIDE_TO = 'AR'
UNK_PAT = '"※ *「*」不在*"'
# 按摘要认的两种 △（「…按摘要认成「x」：对就在收费项选上」「…摘要能对上好几种收费项，按「x」算了：请在收费项里选」）；
# 复核 data-1/F1：别的 △ 里出现「摘要…「」的（冲减提示）不能被它找到
GUESS_PAT = '"△*摘要*「*」*选*"'
DUP_PAT = '"△ 重复粘贴*"'
L_MISS = '※ 这一行没算进来'                                         # 台帐 L 列「读不到的行」的开头（s_input 写的）
NET = '流_净'
F_OKG = Font(name=YH, sz=10, bold=True, color='FF375623')
F_SUMM = Font(name=YH, sz=12, bold=True, color='FFC00000')
F_TITLE_RED = Font(name=YH, sz=15, bold=True, color='FFC00000')
WRAP_TOP = Alignment(horizontal='left', vertical='top', wrap_text=True)
SEC_FILL = 'FFC00000'
NEAR = N_LEDGER // 10                                            # 台帐还能录不到容量的 10%：△ 快满（跟台帐第 1 行红字同一条线）


def _abs(sq, c):
    """表名＋'J4' → 表!$J$4"""
    m = re.fullmatch(r'([A-Z]+)(\d+)', c)
    return f'{sq}!${m.group(1)}${m.group(2)}'


def money(x):
    return f'TEXT({x},"#,##0.00")'


def _f(x):
    """公式文字去掉开头的 ="""
    return x[1:] if isinstance(x, str) and x.startswith('=') else x


# ═════════════════════════════ 直接按口径算（不开窗口） ═════════════════════════════
def s_months(lo, hi, p=None, val=NET, g=None):
    """月码 lo..hi（含）、属性 p、归属公司 g 的 Σval（有效行才有月码）"""
    f = f'SUMIFS({val},流_月码,">="&({lo}),流_月码,"<="&({hi})'
    if p is not None:
        f += f',流_属性号,{p}'
    if g is not None:
        f += f',流_归属号,{g}'
    return f + ')'


def s_year(y, p=None, val=NET, g=None):
    return s_months(f'({y})*100+1', f'({y})*100+12', p, val, g)


def s_days(d1, d2, p=None, val=NET):
    f = f'SUMIFS({val},流_有效,1,流_日期,">="&({d1}),流_日期,"<="&({d2})'
    if p is not None:
        f += f',流_属性号,{p}'
    return f + ')'


def bal_ym(ym):
    """全部账户截至月码 ym（含）的余额"""
    return f'(P_期初合计+SUMIFS(流_净,流_有效,1,流_月码,"<="&({ym})))'


def bal_before_ym(ym):
    return f'(P_期初合计+SUMIFS(流_净,流_有效,1,流_月码,"<"&({ym})))'


def bal_d(d):
    """全部账户截至日期 d（含）的余额"""
    return f'(P_期初合计+SUMIFS(流_净,流_有效,1,流_日期,"<="&({d})))'


# ═════════════════════════════ 核对表 ═════════════════════════════
class Chk:
    def __init__(self, ws):
        self.ws = ws
        self.r = CHK['首行']
        self.no = 0
        self.judge_rows = []           # 有「判断」的行（汇总用）
        self.rows = {}                 # 检查项 → 行号（model/chk 按「检查项」文字找行，这里只给调试用）
        ws._pause_exempt = list(getattr(ws, '_pause_exempt', []) or [])

    def exempt(self, cell):
        self.ws._pause_exempt.append(cell)

    def section(self, text):
        ws, r = self.ws, self.r
        ws.merge_cells(f'A{r}:{C["说明"]}{r}')
        put(ws, f'A{r}', text, F_SEC, fill(SEC_FILL), align=AL)
        for i in range(2, CI(C['说明']) + 1):
            ws.cell(row=r, column=i).border = BD
        ws.row_dimensions[r].height = 20
        self.r += 1
        return r

    def row(self, kind, label, d, judge, note, fmt=ACC, helpers=None, sheet=None, ref='A1', height=None):
        """一项检查：d＝结果公式（不带 =）；judge＝'eq0'（算不平 ✗）/('x'|'w', 单位)（有就 ✗/△）/公式（只准引用本行 $D 和帮助格）；
           note＝说明（文字或 = 开头的公式）；helpers＝[公式…] 依次写进 K、L、M… 列（隐藏）。
           录入模式：结果、说明、帮助格 pause/pause0；判断格登记豁免，显示「— 录入模式中」"""
        ws, r = self.ws, self.r
        self.no += 1
        hc = {}
        for i, h in enumerate(helpers or []):
            col = CL(HELP0 + i)
            ws[f'{col}{r}'] = pause0(h)
            ws[f'{col}{r}'].font = F_HELP
            hc[i] = f'${col}{r}'
        D = f'${C["结果"]}{r}'
        if judge == 'eq0':
            jf = f'IF(ROUND(N({D}),2)=0,"✓ 通过","✗ 算不平")'
        elif isinstance(judge, tuple):
            lvl, unit = judge
            jf = f'IF(N({D})=0,"✓ 没有","{"✗" if lvl == "x" else "△"} 有 "&{D}&" {unit}")'
        else:
            jf = judge
        put(ws, f'{C["序号"]}{r}', self.no, F_AUTO, FILL_NONE, align=AC)
        put(ws, f'{C["类别"]}{r}', kind, F_TXT, FILL_NONE, align=AC)
        c = put(ws, f'{C["检查项"]}{r}', label, F_TXTB, FILL_NONE, align=ALW)
        if sheet:
            c.font = Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single')
            link(c, sheet, ref)
        put(ws, f'{C["结果"]}{r}', pause(d), F_AUTOB, FILL_NONE, fmt=fmt, align=AR)
        put(ws, f'{C["判断"]}{r}', f'=IF({ENTRY},"{J_ENTRY}",{jf})', F_OKG, FILL_NONE, align=AL)
        self.exempt(f'{C["判断"]}{r}')
        put(ws, f'{C["说明"]}{r}', pause(note) if isinstance(note, str) and note.startswith('=') else note, F_TXT, FILL_NONE,
            align=ALW)
        if height:
            ws.row_dimensions[r].height = height
        self.judge_rows.append(r)
        self.rows[label] = r
        self.r += 1
        return r, hc


def _diff_note(static, d_cell, items):
    """说明＝固定文字＋（结果≠0 时）第一个差的项：items＝[(名字, 帮助格)]"""
    ch = '""'
    for nm, cell in reversed(items):
        ch = f'IF({cell}<>0,"{nm}差 "&{money(cell)},{ch})'
    return f'="{static}"&IF(ROUND(N({d_cell}),2)=0,"","　→ "&{ch})'


def _sum_abs(hc, n):
    return 'ROUND(' + '+'.join(f'ABS({hc[i]})' for i in range(n)) + ',2)'


def _multi(ck, kind, label, static, items, sheet):
    """一行比好几项：items＝[(名字, 报表格, 台帐算的公式)]；帮助格＝报表−台帐"""
    hs = [f'ROUND({flow if cell is None else f"N({cell})-({flow})"},2)' for _nm, cell, flow in items]
    r = ck.r
    ref = lambda i: f'${CL(HELP0 + i)}{r}'
    note = _diff_note(static, f'${C["结果"]}{r}', [(nm, ref(i)) for i, (nm, _c, _f) in enumerate(items)])
    return ck.row(kind, label, _sum_abs({i: ref(i) for i in range(len(items))}, len(items)), 'eq0', note, helpers=hs, sheet=sheet)


def _tx(field, k):
    """清单里取台帐原文（_流 位置 k → 台帐那一格）：格子里是错误值（#N/A 这种）时出空，不让整格变成错误值"""
    return f'IFERROR(TRIM({lref_z(field, k)}&""),"")'


def _key_field(k):
    """基础资料某块「有没有这一行」看的那一列（BASE_KEYCOL）对应的字段名"""
    return next(f for f, c in BASE_COLS[k].items() if c == BASE_KEYCOL[k])


def build_chk(ws, ctx):
    ck = Chk(ws)
    widths(ws, {'A': 7, 'B': 9, 'C': 40, 'D': 16, 'E': 24, 'F': 84, 'G': 2})
    tip = ('这张表自动查整本账算得对不对、台帐里还有什么要补要改。判断：✓ 没问题；△ 看一眼、最好补；✗ 要改（勾稽「算不平」多半是报表里的公式被改了，'
           '请告诉我们）。勾稽是拿每张报表上自己选的年份、日期，跟【资金台帐】直接算的比；清单里的「台帐第几行」可以在台帐按 Ctrl+G 输入 A 和行号跳过去。'
           '「台帐分段」一项查插删行有没有把台帐弄乱。基础资料①「录入模式」填了「是」时本表先不算，只留第 3 行的 ※、△ 行数。'
           '全部自动算，不用填；点蓝色的检查项可以跳到那张表。')
    title(ws, '核对表', C['说明'], tip, h2=44)
    ws['A1'] = f'=IF({ENTRY},"{PAUSE_TITLE}","核对表")'
    ws.conditional_formatting.add('A1', FormulaRule(formula=[ENTRY], font=F_TITLE_RED, fill=cf_bg('FFFFD7D7')))
    # ── 计数格（H 列，标签 I 列）──
    gs = '流_项推,">0",流_对方号,0,流_有效,1'
    base_x = '+'.join(f'COUNTIF({bref(k, "核对")},"※*")' for k, *_r in BASE_BLOCKS)
    cnt = {
        '问题': 'COUNTIF(流_问题,"※*")+IFERROR(P_没算+MAX(0,-P_笔数差)+(P_错位>0)+P_锚坏,1)',
        '提醒': 'COUNTIF(流_问题,"△*")+COUNTIFS(流_问题,"",流_配对,1)',
        '按摘要笔数': f'COUNTIFS({gs})',
        '按摘要金额': f'ROUND(SUMIFS(流_收,{gs})+SUMIFS(流_支,{gs}),2)',
        '没配上': 'COUNTIF(流_配对,1)',
        '待分类笔数': f'COUNTIFS(流_有效,1,流_属性号,{I5})',
        '待分类金额': f'ROUND(SUMIFS(流_净,流_有效,1,流_属性号,{I5}),2)',
        '认不出名字': f'COUNTIF(流_问本,{UNK_PAT})',
        '自检': 'N(P_自检)',
        '基础资料问题': base_x,
        '负余额户': f'COUNTIF(${NEG_COL["余额"]}$5:${NEG_COL["余额"]}${4 + N_ACC},"<0")',
    }
    assert set(cnt) == set(CHK_CNT)
    live = ('问题', '提醒')                                    # 录入模式照算（首页提醒条、台帐汇总条、台帐第 1 行要用）
    for k, cell in CHK_CNT.items():
        ws[cell] = '=' + cnt[k] if k in live else pause0(cnt[k])
        ws[cell].font = F_HELP
        if k in live:
            ck.exempt(cell)
        lab = CL(CI(cell[0]) + 1) + cell[1:]
        ws[lab] = '核对_' + k
        ws[lab].font = F_HELP
    # 每个账户的当前余额（块表最后一行：块前＋块净；账户 a 在第 4+a 行）
    nb, ng = NEG_COL['余额'], NEG_COL['负']
    ws[f'{nb}4'] = '账户当前余额'
    for a in range(1, N_ACC + 1):
        ws[f'{nb}{4 + a}'] = pause(f'IF(INDEX(账户_有效,{a})=1,ROUND(INDEX(账户_期初,{a})+INDEX(块_前,{N_BLK},{a})'
                                   f'+INDEX(块_净,{N_BLK},{a}),2),"")')
        ws[f'{ng}{4 + a}'] = pause0(f'IF(AND(ISNUMBER({nb}{4 + a}),N({nb}{4 + a})<0),1,0)' + (f'+{ng}{3 + a}' if a > 1 else ''))
        ws[f'{nb}{4 + a}'].font = ws[f'{ng}{4 + a}'].font = F_HELP
    _seg_helpers(ws)
    # ── 表头 ──
    header(ws, CHK['表头行'], [(C[k], k) for k in ('序号', '类别', '检查项', '结果', '判断', '说明')], '核对')
    ws[f'{C["结果"]}{CHK["表头行"]}'] = '结果\n（差额／个数）'

    # ═══════ 一、勾稽 ═══════
    ck.section('一、勾稽：报表跟【资金台帐】直接算的比（每项用报表自己选的年份、日期；结果是各项差额的绝对值之和，应为 0）')
    L3 = dict((lab, val) for _l, val, lab in LEDGER_KPI)['资金总额']
    blk = lambda nm: f'SUM(INDEX({nm},{N_BLK},1):INDEX({nm},{N_BLK},{N_ACC}))'
    ck.row('勾稽', '资金台帐「资金总额」＝各账户期初＋隐藏的分段累计',
           f'ROUND(N({q(SH_LEDGER)}!${L3[0]}${L3[1:]})-(P_期初合计+{blk("块_前")}+{blk("块_净")}),2)', 'eq0',
           f'台帐第 3 行的资金总额（期初合计＋全部流水收－支）跟隐藏的分段合计（台帐每 {SEG_ROWS} 行一段、各账户累计）比：'
           f'不平说明隐藏的计算或公式被改坏了，请告诉我们', sheet=SH_LEDGER)
    # 收支分类报表
    Y, G = '分类_年份', '分类_公司号'
    fl = lambda p, sg='': f'{sg}IF({G}=0,{s_year(Y, p)},{s_year(Y, p, g=G)})'
    _multi(ck, '勾稽', '收支分类报表：上面 8 个汇总数（全年）', '用报表的年份、公司（空＝全集团，选了公司按归属公司）直接算经营收入、支出、净额、投资建设、往来、待分类、内部、资金净增加',
           [('经营收入', '分类_经营收入', fl(I1)), ('经营支出', '分类_经营支出', fl(I2, '-')),
            ('经营净额', '分类_经营净额', f'{fl(I1)}+{fl(I2)}'), ('投资建设', '分类_投资建设', fl(I3, '-')),
            ('往来净额', '分类_往来净额', fl(I4)), ('待分类', '分类_待分类', fl(I5)), ('内部净额', '分类_内部净额', fl(I9)),
            ('资金净增加', '分类_净增加', f'IF({G}=0,{s_year(Y)},{s_year(Y, g=G)})')], SH_CLS)
    # 收支分项汇总
    lo = 'IF(分项_月码=0,分项_年份*100+1,分项_月码)'
    hi = 'IF(分项_月码=0,分项_年份*100+12,分项_月码)'
    fm = lambda p=None, sg='': f'{sg}{s_months(lo, hi, p)}'
    _multi(ck, '勾稽', '收支分项汇总：合计列 8 个汇总数', '用报表的年份、月份（或全年）直接算全集团的经营收入、支出、净额、投资建设、往来、待分类、内部、资金净增加',
           [('经营收入', '分项_经营收入', fm(I1)), ('经营支出', '分项_经营支出', fm(I2, '-')),
            ('经营净额', '分项_经营净额', f'{fm(I1)}+{fm(I2)}'), ('投资建设', '分项_投资建设', fm(I3, '-')),
            ('往来净额', '分项_往来净额', fm(I4)), ('待分类', '分项_待分类', fm(I5)), ('内部净额', '分项_内部净额', fm(I9)),
            ('资金净增加', '分项_净增加', fm())], SH_ITEMCO)
    r = ck.r
    tip_cls = f'{q(SH_CLS)}!$A${CLS["提示行"]}'
    tip_ic = f'{q(SH_ITEMCO)}!$A${ITEMCO["提示行"]}'
    ck.row('勾稽', '收支分类报表、收支分项汇总：下面明细跟上面汇总',
           f'IFERROR((LEFT({tip_cls},1)="※")*1,0)+IFERROR((LEFT({tip_ic},1)="※")*1,0)',
           f'IF(N($D{r})=0,"✓ 通过","✗ 有 "&$D{r}&" 张报 ※")', '两张表各自的提示行（收支分类第 %d 行、分项汇总第 %d 行）有没有报 ※（明细对不上汇总、收费项太多放不下）'
           % (CLS['提示行'], ITEMCO['提示行']), fmt=INT, sheet=SH_CLS)
    # 收支报表
    Y = '收支_年份'
    _multi(ck, '勾稽', '收支报表：收入总额、支出总额、收支盈亏（全年）', '用报表的年份直接算经营收入（属性收入的收－支）、经营支出（属性支出的支－收）；块 2、块 3 的支出总额都比',
           [('收入总额', '收支_收入总额', s_year(Y, I1)), ('支出总额（按收费项）', '收支_支出总额', '-' + s_year(Y, I2)),
            ('支出总额（按收付人）', '收支_支出总额3', '-' + s_year(Y, I2)), ('收支盈亏', '收支_盈亏', f'{s_year(Y, I1)}+{s_year(Y, I2)}')],
           SH_IE)
    r = ck.r
    neg = [('收入「其余」', '收支_收入其余'), ('支出「其余」', '收支_支出其余'), ('支出「未列示」', '收支_项未列示')]
    hs = [f'IF(ROUND(N({c}),2)<0,1,0)' for _n, c in neg]
    ref = lambda i: f'${CL(HELP0 + i)}{r}'
    nt = '""'
    for i, (nm_, c) in reversed(list(enumerate(neg))):
        nt = f'IF({ref(i)}=1,"{nm_} "&{money(f"N({c})")},{nt})'
    ck.row('提示', '收支报表：「其余」「未列示」不是负数', '+'.join(ref(i) for i in range(3)),
           f'IF(N($D{r})=0,"✓ 没有","△ 有 "&$D{r}&" 项是负数")',
           f'="补差行是负数：前几名里有人退款比收款多，或者收费项、收付人列得比总额多（不算错，看一眼）"&IF(N($D{r})=0,"","　→ "&{nt})',
           fmt=INT, helpers=hs, sheet=SH_IE)
    # 月度汇报表
    Y = '月报_年份'
    yi = lambda p, val=NET, sg='': f'{sg}{s_year(Y, p, val)}'
    _multi(ck, '勾稽', '月度汇报表（一）：全年各列', '用报表的年份直接算经营收入、支出、投资建设、往来流入/流出、待分类、内部转入/转出、资金净增加、年末资金总额',
           [('经营收入', '月报_经营收入', yi(I1)), ('经营支出', '月报_经营支出', yi(I2, sg='-')),
            ('投资建设', '月报_投资建设', yi(I3, sg='-')), ('往来流入', '月报_往来流入', yi(I4, '流_收')),
            ('往来流出', '月报_往来流出', yi(I4, '流_支')), ('待分类', '月报_待分类', yi(I5)), ('内部转入', '月报_内部转入', yi(I9, '流_收')),
            ('内部转出', '月报_内部转出', yi(I9, '流_支')), ('资金净增加', '月报_资金净增加', s_year(Y)),
            ('年末资金', '月报_月末资金', bal_ym(f'{Y}*100+12'))], SH_MON)
    _multi(ck, '恒等式', '月度汇报表（一）：各列加起来＝资金净增加；年末－净增加＝年初',
           '经营收入－经营支出－投资建设＋往来流入－往来流出＋待分类＋内部转入－内部转出＝资金净增加；年末资金－资金净增加＝台帐算的年初资金',
           [('各列之和', '月报_资金净增加', '月报_经营收入-月报_经营支出-月报_投资建设+月报_往来流入-月报_往来流出+月报_待分类'
                                       '+月报_内部转入-月报_内部转出'),
            ('年初资金', '月报_月末资金', f'月报_资金净增加+{bal_before_ym(f"{Y}*100+1")}')], SH_MON)
    _multi(ck, '勾稽', '月度汇报表（二）：指定月份各账户合计', '用二表选的月份直接算：月末余额合计＝全部账户截至月底的余额；收入合计＝当月收入栏（不含内部转账）',
           [('月末余额', '月报2_月末合计', bal_ym('月报2_月码')),
            ('收入', '月报2_收入合计', f'SUMIFS(流_收,流_月码,月报2_月码,流_属性号,"<>{I9}")')], SH_MON, )
    # 资金日报
    _multi(ck, '勾稽', '资金日报（一）：期初、收款、付款、期末合计', '用①的起止日期直接算：期初＝起始日前的余额，收款/付款＝期间收入栏/支出栏（含内部转账），期末＝截止日余额',
           [('期初', '日报_期初合计', f'P_期初合计+SUMIFS(流_净,流_有效,1,流_日期,"<"&日报_起)'),
            ('收款', '日报_收款合计', s_days('日报_起', '日报_止', val='流_收')),
            ('付款', '日报_付款合计', s_days('日报_起', '日报_止', val='流_支')),
            ('期末', '日报_期末合计', bal_d('日报_止'))], SH_DAY)
    _multi(ck, '勾稽', '资金日报（二）：本日余额合计', '用②的日期直接算全部账户那天的余额', [('本日余额', '日报2_余额合计', bal_d('日报2_日'))], SH_DAY)
    _multi(ck, '勾稽', '资金日报（三）：月末余额合计', '用③的月份直接算全部账户那个月底的余额', [('月末余额', '日报3_月末合计', bal_ym('日报3_月码'))],
           SH_DAY)
    # 公司汇总
    Y = '公汇_年份'
    _multi(ck, '勾稽', '公司汇总（一）：合计行（全集团）', '用报表的年份直接算年初资金、经营收入、支出、投资建设、往来、待分类、内部、年末资金',
           [('年初资金', '公汇_年初资金', bal_before_ym(f'{Y}*100+1')), ('经营收入', '公汇_经营收入', s_year(Y, I1)),
            ('经营支出', '公汇_经营支出', '-' + s_year(Y, I2)), ('投资建设', '公汇_投资建设', '-' + s_year(Y, I3)),
            ('往来净额', '公汇_往来净额', s_year(Y, I4)), ('待分类', '公汇_待分类', s_year(Y, I5)),
            ('内部净额', '公汇_内部净额', s_year(Y, I9)), ('年末资金', '公汇_年末资金', bal_ym(f'{Y}*100+12'))], SH_CO)
    ck.row('勾稽', '公司汇总：代收代付差额合计＝0；三表年末合计＝一表', 'ROUND(ABS(N(公汇_代收代付差额))+ABS(N(公汇3_年末)-N(公汇_年末资金)),2)', 'eq0',
           '私户替别的公司收付，按公司看会差一点（代收代付差额），全集团合起来应为 0；三表（公司→账户）年末合计跟一表合计一样', sheet=SH_CO)
    co_rows = f'{q(SH_CO)}!${CO1["列"]["年末资金"]}${CO1["首行"]}:${CO1["列"]["年末资金"]}${CO1["合计行"] - 1}'
    r = ck.r
    ck.row('勾稽', '公司汇总：各公司年末资金加起来＝合计', f'ROUND(SUM({co_rows})-N(公汇_年末资金),2)',
           f'IF(ROUND(N($D{r}),2)=0,"✓ 通过","△ 差 "&{money(f"$D{r}")})',
           '不等于说明有账户的「公司」不在基础资料②（这些账户的钱没归到哪家公司）：去基础资料③核对列看 ※', sheet=SH_CO)
    # 任意时间段
    d1, d2 = 'P_期起', 'P_期止'
    sd = lambda p, sg='': f'{sg}{s_days(d1, d2, p)}'
    sa = lambda val, cond: f'SUMIFS({val},流_有效,1,流_日期,">="&{d1},流_日期,"<="&{d2},流_属性号,{cond})'
    _multi(ck, '勾稽', '任意时间段收支报表：上面 7 个数和各账户合计', '用报表的开始、结束日期直接算经营收支、投资建设、往来、待分类、内部，和各账户期初、收入、支出（不含内部）、内部转入/转出、期末合计',
           [('经营收入', '时段_经营收入', sd(I1)), ('经营支出', '时段_经营支出', sd(I2, '-')),
            ('经营净额', '时段_经营净额', f'{sd(I1)}+{sd(I2)}'), ('投资建设', '时段_投资建设', sd(I3, '-')),
            ('往来净额', '时段_往来净额', sd(I4)), ('待分类', '时段_待分类', sd(I5)), ('内部净额', '时段_内部净额', sd(I9)),
            ('期初余额', '时段_期初余额合计', f'P_期初合计+SUMIFS(流_净,流_有效,1,流_日期,"<"&{d1})'),
            ('收入', '时段_收入合计', sa('流_收', f'"<>{I9}"')), ('支出', '时段_支出合计', sa('流_支', f'"<>{I9}"')),
            ('内部转入', '时段_内部转入合计', sa('流_收', I9)), ('内部转出', '时段_内部转出合计', sa('流_支', I9)),
            ('期末余额', '时段_期末余额合计', f'IF({d2}<{d1},P_期初合计+SUMIFS(流_净,流_有效,1,流_日期,"<"&{d1}),{bal_d(d2)})')], SH_SPAN)
    _multi(ck, '恒等式', '任意时间段：期末合计－期初合计＝各项净额之和', '各账户期末合计－期初合计＝经营净额－投资建设＋往来净额＋待分类＋内部净额',
           [('期末－期初', '时段_期末余额合计', '时段_期初余额合计+时段_经营净额-时段_投资建设+时段_往来净额+时段_待分类+时段_内部净额')], SH_SPAN)
    # 首页、年末余额
    _multi(ck, '勾稽', '首页：资金总额＝截止日全部账户余额', '用首页的截止日期直接算全部账户的余额',
           [('资金总额', '首页_资金总额', bal_d('P_截止'))], SH_HOME)
    _multi(ck, '勾稽', '年末余额：合计＝台帐算的年末资金', '用年末余额表的年份直接算全部账户的年末余额（期初＋到那年 12 月底的收－支）',
           [('年末余额', '年末_合计', bal_ym('年末_年份*100+12'))], SH_YE)
    # 各报表互相对得上（同一年 / 同一月时才比）
    pairs = [
        ('收支分类·经营收入', 'AND(分类_年份=月报_年份,分类_公司号=0)', '分类_经营收入', 'N(月报_经营收入)'),
        ('收支分类·经营支出', 'AND(分类_年份=月报_年份,分类_公司号=0)', '分类_经营支出', 'N(月报_经营支出)'),
        ('收支分类·内部净额', 'AND(分类_年份=月报_年份,分类_公司号=0)', '分类_内部净额', 'N(月报_内部转入)-N(月报_内部转出)'),
        ('公司汇总·经营收入', '公汇_年份=月报_年份', '公汇_经营收入', 'N(月报_经营收入)'),
        ('公司汇总·年末资金', '公汇_年份=月报_年份', '公汇_年末资金', 'N(月报_月末资金)'),
        ('收支报表·收入总额', '收支_年份=月报_年份', '收支_收入总额', 'N(月报_经营收入)'),
        ('收支报表·支出总额', '收支_年份=月报_年份', '收支_支出总额', 'N(月报_经营支出)'),
        ('分项汇总·经营净额', 'AND(分项_月码=0,分项_年份=月报_年份)', '分项_经营净额', 'N(月报_经营收入)-N(月报_经营支出)'),
        ('年末余额·合计', '年末_年份=公汇_年份', '年末_合计', 'N(公汇_年末资金)'),
        ('资金日报③·月末', '日报3_月码=月报2_月码', '日报3_月末合计', 'N(月报2_月末合计)'),
    ]
    _multi(ck, '勾稽', '各报表互相对得上（年份、月份一样时才比）',
           '同一年：收支分类（全集团）、公司汇总、收支报表、分项汇总（看全年时）、年末余额都跟月度汇报表比；同一个月：资金日报③月末跟月度汇报表二比',
           [(nm_, None, f'IF({cond},N({a})-({b}),0)') for nm_, cond, a, b in pairs], SH_MON)
    # 内部转账、自检
    ck.row('勾稽', '内部转账：本年净额＝0（管理年度）', f'ROUND({s_year("P_管理年度", I9)},2)', 'eq0',
           '本册两个账户之间互转，转出、转入两头都记了、金额一样，全集团净额就是 0。不是 0：有的只记了一头或金额不对，看下面「六、内部转账没配上」',
           sheet=SH_LEDGER, ref='A4')
    r = ck.r
    ck.row('规则', '「按摘要认」规则推断自检', 'N(P_自检)', f'IF(N($D{r})=0,"✓ 通过","✗ 不一致 "&$D{r}&" 个")',
           '用几条写死的规则和摘要试算「按摘要认收费项」：不一致说明这个软件（版本）算不了按摘要认，请告诉我们用的是什么软件、哪个版本',
           fmt=INT, sheet=SH_BASE)

    # ═══════ 二、子表 ═══════
    ck.section(f'二、子表（{len(SUB_SHEETS)} 张账户子表＋备用 {len(SPARE_SHEETS)} 张＋账户查询）：余额、月份表合计、笔数、明细收－支跟【资金台帐】直接算的比')
    subs = [(s, o, SUB_CAP[s], 'sub') for s, o, _c in SUB_SHEETS] + [(s, '', SUB_CAP[s], 'spare') for s, _c in SPARE_SHEETS] + \
           [(SH_QRY, '', N_QUERY, 'qry')]
    P = SUB_PARAM
    for s, orig, cap, kind in subs:
        _sub_row(ck, s, orig, cap, kind, P)

    # ═══════ 三、待办 ═══════
    ck.section('三、待办：台帐里要补、要改的（「约」＝数的是这一行显示的第一个提示，被排在前面的提示盖住的数不到）')
    ck.row('待办', '台帐有 ※ 的行（必须改）', CNT['问题'], ('x', '行'),
           '台帐 L 列红底的行：日期、账户、金额认不出，收费项不在清单，内部转账没填另一户……这些行不算或算错，必须改。前 50 行见下面「七」；'
           '有行没算进来、格子被挪乱的也算在这里，看下面「台帐分段」', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '台帐有 △ 的行（最好补）', CNT['提醒'], ('w', '行'),
           '台帐 L 列琥珀色字的行：照算了，但最好补齐（按摘要认的收费项、收付人没填、内部转账没配上、日期顺序……）', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '名字认不出（改了名没加别名？）', CNT['认不出名字'], ('x', '行'),
           '账户、收费项、转账的另一户、归属公司写的名字在基础资料里找不到：多半是改了名没把旧名加进「别名」。各种写法见下面「五」', fmt=INT,
           sheet=SH_BASE, ref='A1')
    ck.row('待办', '收费项按摘要认的（建议固化）', CNT['按摘要笔数'], ('w', '笔'),
           f'="共 "&{money(CNT["按摘要金额"])}&" 元（收＋支）。收费项空着、按基础资料⑥规则猜的（台帐 N 列琥珀色斜体）：核对对了就把 N 列粘成数值到 C 列（使用说明「五」）；前 50 笔见下面「八」"',
           fmt=INT, sheet=SH_LEDGER, ref='A4')
    # 复核 F11：写清是净额、在台帐 M 列筛「待分类」就是这几行
    ck.row('待办', '待分类（收费项待定、没填又猜不出）', CNT['待分类笔数'], ('w', '笔'),
           f'="净额 "&{money(CNT["待分类金额"])}&" 元（收、付分开的数见数据整理记录「待确认」）。不算经营收支。是哪几行：在【资金台帐】M 列（属性）'
           f'点筛选箭头只勾「待分类」；改 C 列收费项，或在基础资料④给收费项定属性"', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '内部转账没配上', CNT['没配上'], ('w', '笔'),
           f'同一对账户前后 {PAIR_DAYS} 天内找不到金额相反的一笔：漏录一头、金额或日期对不上。前 50 笔见下面「六」', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '收付人不在基础资料⑤（约）', 'COUNTIF(流_提本,"△ 收付人「*」不在⑤*")', ('w', '笔'),
           '台帐填了收付人但⑤里没有：报表里算「没填」。去⑤加上，或把旧写法加进别名', fmt=INT, sheet=SH_BASE)
    ck.row('待办', '收付人没填（收入、支出类）（约）', 'COUNTIF(流_提本,"△ 收付人没填*")', ('w', '笔'),
           '经营收入、支出的流水没填收付人，摘要也猜不出：收支报表里算「没填」', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '本账户日期顺序不对（约）', 'COUNTIF(流_顺序,"△ 本账户上面有日期更晚的*")', ('w', '行'),
           '同一个账户里，上面有日期更晚的：余额（H 列）按上下顺序算，点「日期」列的筛选箭头升序排一下就好', fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '可能重复粘贴', f'COUNTIF(流_提本,{DUP_PAT})', ('w', '行'),
           f'同账户、同一天、同金额、同摘要（备注、原表位置也一样），前后 {DUP_WIN} 行里还有一笔：粘了两遍？真是两笔就在备注写一句。见下面「九」',
           fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '日期跳期（年份写错了？）', 'COUNTIF(流_顺序,"※ 比前面最晚的一笔*")', ('x', '笔'),
           f'比前面最晚的一笔晚 {JUMP_DAYS} 天以上、下一笔又早回去（比如 10 月的账里夹了一笔敲成 12 月的）：钱照算，但不算进「最后日期」。改对年份、月份', fmt=INT,
           sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '最后几笔日期一下跳得很远', 'COUNTIF(流_顺序,"△ 比前面最晚的一笔晚*")', ('w', '笔'),
           f'台帐最后几笔比前面最晚的一笔晚 {JUMP_DAYS} 天以上，后面还没有别的流水：日期没写错就不用管，接着录正常日期的流水后这条自己消失', fmt=INT,
           sheet=SH_LEDGER, ref='A4')
    ck.row('待办', '暂停/销户的账户还在记（约）', 'COUNTIF(流_提本,"△ 账户已暂停*")', ('w', '笔'),
           '基础资料③状态是暂停/销户、停用日以后还有流水：真在用就把状态改回「在用」', fmt=INT, sheet=SH_BASE)
    r = ck.r
    nb_rng = f'${NEG_COL["余额"]}$5:${NEG_COL["余额"]}${4 + N_ACC}'
    first_neg = f'IFERROR(MATCH(1,${NEG_COL["负"]}$5:${NEG_COL["负"]}${4 + N_ACC},0),1)'
    ck.row('待办', '当前余额是负数的账户', CNT['负余额户'], ('w', '户'),
           f'="漏录了收入、或者期初不对？"&IF(N($D{r})=0,"","　→ 比如 "&INDEX(账户_名称,{first_neg})&"："&{money(f"INDEX({nb_rng},{first_neg})")})',
           fmt=INT, sheet=SH_HOME)
    base_w = '+'.join(f'COUNTIF({bref(k, "核对")},"△*")' for k, *_r in BASE_BLOCKS)
    ck.row('待办', '基础资料核对列有 ※', CNT['基础资料问题'], ('x', '处'),
           f'="基础资料各块最右边「核对」列：重名、公司不在②、属性填错、规则的收费项认不出……另有 △ "&({base_w})&" 处看一眼"', fmt=INT, sheet=SH_BASE)

    # ═══════ 四、容量和结构 ═══════
    ck.section('四、容量和结构')
    r = ck.r
    ck.row('容量', '资金台帐已用行数', 'P_已用行数',
           f'IF(P_可录<=0,"✗ 满了",IF(P_可录<{NEAR},"△ 快满了","✓ 够用"))',
           '="一年一本：台帐能录 "&TEXT(P_台帐容量,"#,##0")&" 行，已用 "&TEXT(N($D' + str(r) + '),"#,##0")&" 行，还能录约 "'
           '&TEXT(P_可录,"#,##0")&" 行。满了年底按使用说明「八」开新本；年中满了找我们放大（空行别删：删了能录的行就少了）"',
           fmt=INT, sheet=SH_LEDGER, ref='A4')
    ck.row('容量', '已经有下一年的流水（该开新本）', 'COUNTIFS(流_有效,1,流_月码,">="&(P_开账年+1)*100+1)', ('w', '笔'),
           '="本本 "&P_开账年&" 年（次年 1 月的先记在这本也行）：到了 1 月按使用说明「八」开新本"', fmt=INT, sheet=SH_YE)
    ck.row('容量', '开新本后旧年流水没删', 'IF(COUNTIF(流_有效,1)=0,0,COUNTIFS(流_日期,">0",流_日期,"<"&P_开账日))', ('w', '笔'),
           '="开账日（"&TEXT(P_开账日,"yyyy/m/d")&"）以前的流水不算：开新本时第 ④ 步要把上一年的行清掉（使用说明「八」）"', fmt=INT,
           sheet=SH_LEDGER, ref='A4')
    _seg_row(ck)
    r = ck.r
    ncol = CI(LEDGER_LAST)
    ck.row('结构', '资金台帐的列数（插了列/删了列）', f'COLUMNS({q(SH_LEDGER)}!$A${HDR}:${LEDGER_LAST}${HDR})-{ncol}',
           f'IF(N($D{r})<0,"✗ 删了列",IF(N($D{r})>0,"△ 插了列","✓ 没变"))',
           f'台帐 A～{LEDGER_LAST} 共 {ncol} 列，别插列、别删列：删了列整本算不出来，马上撤销（Ctrl+Z）；插了列照样算，但最好删掉。'
           f'右边 {LEDGER_ANCHOR} 列（勿删）也别删', fmt=INT, sheet=SH_LEDGER, ref='A4')
    _disp_row(ck)
    _base_ins_row(ck)
    r = ck.r
    hs, items = [], []
    for k, _t, _tr, hdr, cap in BASE_BLOCKS:
        cols = [CI(c) for f, c in BASE_COLS[k].items() if f not in ('序号', '核对')]
        c1, c2 = CL(min(cols)), CL(max(cols))
        area = f'{q(SH_BASE)}!${c1}${hdr}:${c2}${BASE_LAST[k]}'
        hs.append(f'COUNTA({area})-COUNTA(INDEX({area},1,1):INDEX({area},{cap + 1},COLUMNS({area})))')
        items.append(k)
    ref = lambda i: f'${CL(HELP0 + i)}{r}'
    nt = '""'
    for i, k in reversed(list(enumerate(items))):
        nt = f'IF({ref(i)}>0,"{k}块容量 {BASE_CAP[k]} 个以后有 "&{ref(i)}&" 格",{nt})'
    ck.row('容量', '基础资料各块容量以外有字', '+'.join(ref(i) for i in range(len(items))), ('x', '格'),
           f'="各块只认前「容量」行（公司 {N_CO}、账户 {N_ACC}、收费项 {N_ITEM}、收付人 {N_PARTY}、规则 {N_RULE}），容量以外的不算"&IF(N($D{r})=0,"","　→ "&{nt})',
           fmt=INT, helpers=hs, sheet=SH_BASE)
    ck.row('容量', f'推收付人的规则超过 {N_RULE_P} 条', f'MAX(0,MAX(规_人计数)-{N_RULE_P})', ('x', '条'),
           f'="摘要关键词规则有效 "&P_规则数&" 条：推收费项 "&P_项规则数&" 条、推收付人 "&MAX(规_人计数)&" 条（推收付人的最多 {N_RULE_P} 条，多的不用）"',
           fmt=INT, sheet=SH_BASE)
    ck.row('容量', f'在「收付人」里找的收费项规则超过 {N_RULE_Q} 条', f'MAX(0,N(P_Q规则数)-{N_RULE_Q})', ('x', '条'),
           f'="推收费项的规则里，查找位置空着或填「收付人」的有 "&N(P_Q规则数)&" 条（最多 {N_RULE_Q} 条）：超了以后按收付人认收费项全都不起作用'
           f'（按摘要照样认）。只在摘要里找的规则，把⑥「查找位置」填「摘要」"', fmt=INT, sheet=SH_BASE)
    kw = bref('规则', '关键词')
    ck.row('规则', '摘要关键词规则没起作用的', f'MAX(0,COUNTA(INDEX({kw},2):INDEX({kw},{N_RULE + 1}))-P_规则数)', ('w', '条'),
           '基础资料⑥填了关键词、但收费项和收付人都认不出（或方向、查找位置写错）：这条规则不起作用，看⑥的核对列', fmt=INT, sheet=SH_BASE)
    ck.row('容量', f'收费项大类超过 {N_GROUP} 个', f'MAX(0,N(大类_数)-{N_GROUP})', ('x', '个'),
           f'="大类现在 "&N(大类_数)&" 个：超过 {N_GROUP} 个的报表不单列（归到「大类空或超过」），合并几个大类"', fmt=INT, sheet=SH_BASE)
    ck.row('容量', '收支分类报表、分项汇总放不下', f'MAX(0,COUNT(分类_候选键)-{N_CLS})', ('x', '行'),
           f'="两张表各 {N_CLS} 行（段标题＋大类小计＋收费项×属性），现在要 "&COUNT(分类_候选键)&" 行：多的列不出来（汇总数照样对）"', fmt=INT,
           sheet=SH_CLS)
    r = ck.r
    caps = [('公司数', N_CO), ('账户数', N_ACC), ('收费项数', N_ITEM), ('收付人数', N_PARTY), ('规则数', N_RULE)]
    ck.row('容量', '名单用了多少', 'MAX(' + ','.join(f'N(P_{k})/{n}' for k, n in caps) + ')',
           f'IF(N($D{r})>=0.9,"△ 快满了","✓ 够用")',
           '="公司 "&P_公司数&"/' + str(N_CO) + '、账户 "&P_账户数&"/' + str(N_ACC) + '、收费项 "&P_收费项数&"/' + str(N_ITEM) +
           '、收付人 "&P_收付人数&"/' + str(N_PARTY) + '、规则 "&P_规则数&"/' + str(N_RULE) + '（结果是用得最满的那一块）"', fmt=PCT, sheet=SH_BASE)
    last = ck.r - 1
    # ── 清单 ──
    ck.r += 1
    _unk_list(ck)
    _list(ck, '六、内部转账没配上', CNT['没配上'], '流_配对', '1', '对方｜摘要',
          lambda k: f'"对方："&{_tx("对方账户", k)}&"｜"&{_tx("摘要", k)}')
    _list(ck, '七、※ 的行', 'COUNTIF(流_问题,"※*")', '流_问题', '"※*"', '问题｜摘要',
          lambda k: f'INDEX(流_问题,{k})&"｜"&{_tx("摘要", k)}')
    _list(ck, '八、按摘要认的收费项', CNT['按摘要笔数'], '流_提本', GUESS_PAT, '认成｜摘要（被「重复粘贴」「待定」提示盖住的不在这里）',
          lambda k: f'"认成「"&INDEX(项_名称,INDEX(流_项号,{k}))&"」｜"&{_tx("摘要", k)}')
    _list(ck, '九、可能重复粘贴', f'COUNTIF(流_提本,{DUP_PAT})', '流_提本', DUP_PAT, '摘要｜备注',
          lambda k: f'{_tx("摘要", k)}&"｜"&{_tx("备注", k)}')
    # ── 第 3 行汇总（录入模式照算：只看计数格 H1/H2 和判断列）──
    E = f'${C["判断"]}${CHK["首行"]}:${C["判断"]}${last}'
    home_link(ws, 'A3')
    put(ws, 'B3', '结果', F_KPI_L, FILL_KPI, align=AC)
    ws.merge_cells('C3:E3')
    put(ws, 'C3', f'=IF({ENTRY},"录入模式中：核对表先不算（台帐的 ※、△ 行数照算，见右边）",'
                  f'"✗ "&COUNTIF({E},"✗*")&" 项　　△ "&COUNTIF({E},"△*")&" 项　　✓ "&COUNTIF({E},"✓*")&" 项（共 "&'
                  f'(COUNTIF({E},"✗*")+COUNTIF({E},"△*")+COUNTIF({E},"✓*"))&" 项）")', F_SUMM, FILL_KPI, align=AC)
    for c in ('D3', 'E3'):
        ws[c].fill, ws[c].border = FILL_KPI, BD
    put(ws, 'F3', f'="资金台帐：※ "&{CNT["问题"]}&" 行要改，△ "&{CNT["提醒"]}&" 行最好补（L 列"&IF({ENTRY},"）","；清单在下面）")',
        F_KPI_L, FILL_KPI, align=AL)
    ck.exempt('C3')
    ck.exempt('F3')
    ws.row_dimensions[3].height = 26
    # 条件格式：✗ 红、△ 黄底、— 灰
    rng = f'{C["判断"]}{CHK["首行"]}:{C["判断"]}{last}'
    e0 = f'{C["判断"]}{CHK["首行"]}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({e0},1)="✗"'], stopIfTrue=True,
                                                   font=Font(name=YH, bold=True, color='FFC00000'), fill=cf_bg('FFFFD7D7')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({e0},1)="△"'], stopIfTrue=True,
                                                   font=Font(name=YH, bold=True, color='FF8B5E00'), fill=cf_bg('FFFFF2CC')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({e0},1)="—"'], stopIfTrue=True,
                                                   font=Font(name=YH, color='FF808080')))
    ws.conditional_formatting.add('C3', FormulaRule(formula=[f'AND(NOT({ENTRY}),COUNTIF({E},"✗*")=0)'],
                                                    font=Font(name=YH, sz=12, bold=True, color='FF375623')))
    ws.auto_filter.ref = f'A{CHK["表头行"]}:{C["说明"]}{last}'
    ws.freeze_panes = f'A{CHK["首行"]}'
    hcols = ['H', 'I'] + [CL(i) for i in range(HELP0, CI(HIDE_TO) + 1)]
    for c in hcols:
        ws.column_dimensions[c].width = 10
    hide(ws, *hcols)
    print_setup(ws, f'{CHK["表头行"]}:{CHK["表头行"]}')
    ws._chk_rows = ck.rows
    return ck


# ───────────── 二、子表一行 ─────────────
def _sub_row(ck, s, orig, cap, kind, P):
    """帮助格（K 列起）：0 活 1 账户号 2 年份 3 月码 4 窗首 5 窗末（那一年/那个月）6 年底窗末 7 月底窗末 8 笔数 9 年末余额 10 期末余额
       11 台帐收－支 12 明细收－支 13 差 14 J4 15 台帐收入栏 16 内部转入 17 内部转出（本年或选的月）。
       v3.1：只扫窗口（win_first/win_last：月码在那一年／那个月的行一定在窗口里；余额扫 [1, 年底／月底窗口末]），结果跟整列一样"""
    r = ck.r
    sq = q(s)
    h = lambda i: f'${CL(HELP0 + i)}{r}'
    cell = lambda c: _abs(sq, c)
    end = SUB_R0 + cap - 1
    G_, H_ = SUB_COLS['收入'], SUB_COLS['支出']
    a, y, ym, wf, wl, wye, wme = h(1), h(2), h(3), h(4), h(5), h(6), h(7)
    j4, k4, r19 = cell(SUB_KPI['余额']), cell(SUB_KPI['笔数']), cell(f'{SUB_MON["月末余额"]}{SUB_MON_TOT}')
    mt = {k: cell(f'{SUB_MON[k]}{SUB_MON_TOT}') for k in ('收入', '支出', '内部转入', '内部转出')}   # 月份表合计行（第 19 行）
    w = lambda nm: win(nm, wf, wl)
    kin = lambda val, lo_m, hi_m: (f'SUMIFS({w(val)},{w("流_账年属月键")},">="&{k_aypm(a, y, I9, lo_m)},'
                                   f'{w("流_账年属月键")},"<="&{k_aypm(a, y, I9, hi_m)})')
    mm = f'MOD({ym},100)'
    per = lambda val: (f'IF({ym}=0,SUMIF({w("流_账年键")},{a}*10000+{y},{w(val)}),'
                       f'SUMIF({w("流_账月键")},{a}*{K_YM}+{ym},{w(val)}))')
    hs = [
        f'IF(ISERROR({cell(P["账户号"])}),0,1)',
        f'IFERROR(N({cell(P["账户号"])}),0)',
        f'IFERROR(N({cell(P["年份"])}),0)',
        f'IFERROR(N({cell(P["月码"])}),0)',
        f'IF({ym}=0,{win_first(f"{y}*100+1")},{win_first(ym)})',
        f'IF({ym}=0,{win_last(f"{y}*100+12", wf)},{win_last(ym, wf)})',
        win_last(f'{y}*100+12'),
        f'IF({ym}=0,{wye},{win_last(ym)})',
        (f'IF({a}=0,0,IF({ym}=0,COUNTIF({w("流_账年键")},{a}*10000+{y}),COUNTIF({w("流_账月键")},{a}*{K_YM}+{ym})))'),
        (f'IF({a}=0,0,ROUND(INDEX(账户_期初,{a})+SUMIFS({win("流_净", 1, wye)},{win("流_账年键", 1, wye)},">="&{a}*10000,'
         f'{win("流_账年键", 1, wye)},"<="&{a}*10000+{y}),2))'),
        (f'IF({a}=0,0,IF({ym}=0,{h(9)},ROUND(INDEX(账户_期初,{a})+SUMIFS({win("流_净", 1, wme)},{win("流_账月键", 1, wme)},'
         f'">="&{a}*{K_YM},{win("流_账月键", 1, wme)},"<="&{a}*{K_YM}+{ym}),2)))'),
        f'IF({a}=0,0,ROUND({per("流_净")},2))',
        f'IFERROR(ROUND(SUM({sq}!${G_}${SUB_R0}:${G_}${end})-SUM({sq}!${H_}${SUB_R0}:${H_}${end}),2),0)',
        (f'IF({a}=0,0,IFERROR(ROUND(ABS(N({j4})-{h(9)})+ABS(N({r19})-{h(10)})+ABS(N({k4})-{h(8)})'
         f'+IF({h(8)}<={cap},ABS({h(12)}-{h(11)}),0)+ABS(N({mt["收入"]})-({h(15)}-{h(16)}))+ABS(N({mt["支出"]})-({h(15)}-{h(11)}-{h(17)}))'
         f'+ABS(N({mt["内部转入"]})-{h(16)})+ABS(N({mt["内部转出"]})-{h(17)}),2),-1))'),
        f'IFERROR(N({j4}),0)',
        f'IF({a}=0,0,ROUND({per("流_收")},2))',
        f'IF({a}=0,0,ROUND(IF({ym}=0,{kin("流_收", 1, 12)},{kin("流_收", mm, mm)}),2))',
        f'IF({a}=0,0,ROUND(IF({ym}=0,{kin("流_支", 1, 12)},{kin("流_支", mm, mm)}),2))',
    ]
    named = kind == 'sub'
    lbl = (f'子表【{s}】' + (f'（原表：{orig}）' if orig and orig != s else '')) if kind != 'qry' else '【账户查询】'
    no_acc = '"△ 认不出账户"' if named else '"— 没选账户"'
    judge = (f'IF({h(0)}=0,"✗ 子表没了",IF({a}=0,{no_acc},'
             f'IF({h(8)}>{cap},"✗ 放不下："&{h(8)}&" 笔＞容量 {cap:,}",IF({h(13)}<>0,"✗ 跟台帐对不上","✓ 通过"))))')
    none_txt = ('基础资料③找不到子表名是「' + s + '」的账户：③的「子表名」那一列别改' if named
                else '在 D3 选了账户才有数（给新开的卡用；表名别改）' if kind == 'spare' else '在 D3 选了账户才有数')
    mtxt = f'IF({ym}>0," "&MOD({ym},100)&" 月","")' if kind == 'qry' else '""'
    # 复核 F4：改表名公式会跟着改，不会报这一句；只有删了才报
    note = (f'=IF({h(0)}=0,"这张表被删了：子表别删、别改名（不用的可以隐藏）",IF({a}=0,"{none_txt}",'
            f'{y}&" 年"&{mtxt}&"：台帐 "&{h(8)}&" 笔（本表放得下 {cap:,} 笔），余额 "&{money(h(9))}&'
            f'IF({h(13)}<>0,"；子表余额 "&{money(h(14))}&"、明细收－支 "&{money(h(12))}&"、台帐收－支 "&{money(h(11))},"")))')
    ck.row('子表', lbl, f'IF({h(0)}=0,"",IF({a}=0,"",{h(13)}))', judge, note, helpers=hs, sheet=s, ref='A1')


# ───────────── 四、台帐分段（设计 §15.1.5） ─────────────
def _seg_helpers(ws):
    """隐藏列（第 4+k 行＝第 k 段）：
       AN 错位累计：P_错位>0 时，本表自己按 台k_字段 的首行、行数跟 段_区首/区末 比（段表叶子 段_错位 只准 _参 引用），第一个错位的段＝MATCH(1,AN)
       AO 没算累计：P_没算>0 时，漏行>0 或 夹缝>0 的段累计（只看段表的链，改数据不脏）
       AQ5..AQ9 前 5 处的段号；AP5..AP9 那几处「第 a～b 行」；AR5 容量外（P_台帐末 以下）有字的行数"""
    X, Cc, Pp, Kk, O = (SEG_COL[k] for k in ('错位', '没算', '位置', '段', '容量外'))
    for c, t in ((X, '错位累计'), (Cc, '没算累计'), (Pp, '读不到的几处'), (Kk, '段号'), (O, '容量外行数')):
        ws[f'{c}4'] = t
        ws[f'{c}4'].font = F_HELP
    for k in range(1, N_SEG + 1):
        r = 4 + k
        prev = f'{X}{r - 1}+' if k > 1 else ''
        terms = '+'.join(f'(ROW(INDEX({seg_name(f, k)},1))<>INDEX(段_区首,{k}))'
                         f'+(ROWS({seg_name(f, k)})<>INDEX(段_区末,{k})-INDEX(段_区首,{k})+1)' for f in LEDGER_COLS)
        ws[f'{X}{r}'] = pause0(f'IF(P_错位=0,0,{prev}IF(IFERROR({terms},1)>0,1,0))')
        prev = f'{Cc}{r - 1}+' if k > 1 else ''
        ws[f'{Cc}{r}'] = pause0(f'IF(P_没算=0,0,{prev}IF(IFERROR((INDEX(段_漏行,{k})>0)+(INDEX(段_夹缝,{k})>0),0)>0,1,0))')
        ws[f'{X}{r}'].font = ws[f'{Cc}{r}'].font = F_HELP
    crng = f'${Cc}$5:${Cc}${4 + N_SEG}'
    for i in range(1, 6):
        r = 4 + i
        k = f'{Kk}{r}'
        ws[k] = pause0(f'IFERROR(MATCH({i},{crng},0),0)')
        pt = f'IF({k}=1,INDEX(段_首,1),INDEX(段_尾,{k}-1))'
        ws[f'{Pp}{r}'] = pause(
            f'IF({k}=0,"",IFERROR(IF(INDEX(段_漏行,{k})>0,"第 "&(INDEX(段_前末,{k})+1)&"～"&MIN(INDEX(段_读到,{k}),INDEX(段_区首,{k})-1)&" 行",'
            f'"第 "&({pt}+1)&"～"&MIN(INDEX(段_尾,{k}),INDEX(段_区首,{k})-1)&" 行"),""))')
        ws[k].font = ws[f'{Pp}{r}'].font = F_HELP
    i0 = 'P_台帐末-INDEX(段_首,1)+2'
    cnt = lambda f: f'COUNTA(INDEX(台帐_{f},{i0}):INDEX(台帐_{f},ROWS(台帐_{f})))'
    ws[f'{O}5'] = pause0(f'IF(P_没算=0,0,IFERROR(MAX({cnt("日期")},{cnt("账户")}),0))')
    ws[f'{O}5'].font = F_HELP


def _seg_row(ck):
    """「台帐分段」一行：帮助格 0＝状态（4 锚坏 3 错位/读两遍 2 没算 1 整段删掉 0 正常）1＝第一个错位的段 2＝第一个整段删掉的段
       3＝第一个借位的段 4＝前 5 处读不到的位置（文字）"""
    r = ck.r
    h = lambda i: f'${CL(HELP0 + i)}{r}'
    X, Pp, O = SEG_COL['错位'], SEG_COL['位置'], SEG_COL['容量外']
    # 前 5 处是连着的（第 i 处有，前面几处一定有）：第 1 处不加顿号；容量外接在后面
    places = f'${Pp}$5&' + '&'.join(f'IF(${Pp}${4 + i}="","","、"&${Pp}${4 + i})' for i in range(2, 6))
    far = f'IF(${O}$5>0,IF(${Pp}$5="","","、")&"第 "&(P_台帐末+1)&" 行以下（容量外 "&${O}$5&" 行）","")'
    hs = [
        'IFERROR(IF(P_锚坏=1,4,IF(OR(P_错位>0,P_笔数差<0),3,IF(P_没算>0,2,IF(COUNTIF(段_状态,3)>0,1,0)))),4)',
        f'IFERROR(MATCH(1,${X}$5:${X}${4 + N_SEG},0),0)',
        'IFERROR(MATCH(3,段_状态,0),0)',
        'IFERROR(MATCH(1,段_状态,0),0)',
        f'IFERROR(MID({places}&{far},1,600),"")',
    ]
    st = h(0)
    D = f'$D{r}'
    k1, k3, kb = h(1), h(2), h(3)
    seg_rows = lambda k: f'IFERROR("（台帐第 "&INDEX(段_首,{k})&"～"&INDEX(段_末,{k})&" 行）","")'
    msg = {
        4: '"台帐的表头（第 4 行）或右边的 Q 列（勿删）被删了：所有数都没算。马上撤销（Ctrl+Z）；已经存盘了找我们（重出一本模板，把数导过去）"',
        3: (f'"台帐被插/删了单元格（不是整行），或者只剪了 A～P 的格子、删了中间的列、用了「删除重复项」"&IF({k1}>0,"：从第 "&{k1}&" 段"&{seg_rows(k1)}&"起",'
            f'"")&"后面的数全乱了。马上撤销（Ctrl+Z）；插行、删行一律点左边行号选整行。已经存盘了：按原样把那几格挪回去，或者找我们"'),
        2: (f'"台帐有 "&{D}&" 行没算进来："&IF({h(4)}="","（位置看不出来）",{h(4)})&"。多半是一段里一次插了 {SEG_GROW_MAX} 行以上，'
            f'或者删掉了整段（一次删了 {SEG_ROWS} 多行）又在接缝插了行，或者录到了容量以外。改法：把这几行复制 → 粘到台帐最下面的空行 → '
            f'点行号删掉原来那几整行 → 按日期排序；录到容量以外的：年底开新本，年中满了找我们放大"'),
        1: (f'"第 "&{k3}&" 段整段删掉了（一次删了 {SEG_ROWS} 多行）"&{seg_rows(k3)}&"：数没错，台帐还能录约 "&P_可录&" 行。'
            f'以后大批删行先选中按 Delete 清内容，空行留着"'),
        0: (f'IF({kb}>0,"正常。第 "&{kb}&" 段多插了 "&IFERROR(INDEX(段_尾,{kb})-INDEX(段_读到,{kb}),"")&" 行，下一段接着读（正常）",'
            f'"正常：台帐每 {SEG_ROWS} 行一段读进隐藏的计算表；点行号插删整行、排序、粘贴都不影响。这一项查表头和 Q 列有没有被删、'
            f'格子有没有被挪乱、有没有行没算进来")'),
    }
    note = '=IFERROR(CHOOSE(' + st + '+1,' + ','.join(msg[i] for i in range(5)) + '),"（算不出来：多半是 Q 列或表头被删了，马上撤销）")'
    judge = (f'CHOOSE(N({st})+1,"✓ 正常","△ 有整段删掉了","✗ 有 "&{D}&" 行没算进来","✗ 格子被挪乱","✗ 表头或 Q 列被删")')
    ck.row('结构', '台帐分段（插删行有没有把台帐弄乱）', 'IFERROR(P_没算+MAX(0,-P_笔数差)+(P_错位>0)+P_锚坏,1)', judge, note,
           fmt=INT, helpers=hs, sheet=SH_LEDGER, ref='A4', height=48)


def _disp_row(ck):
    """显示列核对（设计 §15.1.6）：只当「缺公式、被盖掉」用（△）；错位归「台帐分段」。
       帮助格 0 余额个数差 1 余额合计不等 2 ※ 行数不等 3 △ 行数不等"""
    r = ck.r
    h = lambda i: f'${CL(HELP0 + i)}{r}'
    bal, prob = lrange('余额'), lrange('问题')
    hs = [
        f'ABS(COUNT({bal})-COUNTIF(流_有效,1))',
        f'IF(ROUND(SUM({bal})-SUMIF(流_有效,1,流_余额),2)<>0,1,0)',
        f'IF(COUNTIF({prob},"※*")-COUNTIF({prob},"{L_MISS}*")<>COUNTIF(流_问题,"※*"),1,0)',
        f'IF(COUNTIF({prob},"△*")<>{CNT["提醒"]},1,0)',
    ]
    ck.row('结构', '资金台帐显示列缺公式或被盖掉', '+'.join(h(i) for i in range(4)), ('w', '处'),
           f'="台帐 {LEDGER_SHOW["余额"]} 列（余额）、{LEDGER_SHOW["问题"]} 列（问题）跟自动算的对不上"&IF({h(0)}>0,"（有余额的行差 "&{h(0)}&" 行）","")'
           f'&"：多半是中间插了行、新行没带公式，或者被粘贴盖掉了。选中上一行的 {LEDGER_SHOW["余额"]}、{LEDGER_SHOW["问题"]}～{LEDGER_SHOW["收付人认"]} '
           f'往下拖一格（使用说明「十」）"', fmt=INT, helpers=hs, sheet=SH_LEDGER, ref='A4')


def _base_ins_row(ck):
    """复核 compat-3：基础资料中间插的行没有序号、核对公式（插出来的行不带公式）→ 各块「有名字但序号空」的行数"""
    r = ck.r
    h = lambda i: f'${CL(HELP0 + i)}{r}'
    hs, names = [], []
    for k, t, _tr, _hdr, cap in BASE_BLOCKS:
        key, no = bref(k, _key_field(k)), bref(k, '序号')
        hs.append(f'COUNTIFS(INDEX({key},2):INDEX({key},{cap + 1}),"<>",INDEX({no},2):INDEX({no},{cap + 1}),"")')
        names.append(t.split('（')[0].replace(' ', ''))
    nt = '""'
    for i, nm in reversed(list(enumerate(names))):
        nt = f'IF({h(i)}>0,"{nm}有 "&{h(i)}&" 行",{nt})'
    chk_cols = '、'.join(f'{t.split("（")[0].split(" ")[0]}{BASE_COLS[k]["核对"]}' for k, t, *_r in BASE_BLOCKS)
    ck.row('结构', '基础资料中间插的行没带公式', '+'.join(h(i) for i in range(len(hs))), ('x', '行'),
           f'="中间插出来的行没有序号和核对公式，这些行填错了（重名、期初没填、规则的收费项认不出……）就查不出来"&IF(N($D{r})=0,"","　→ "&{nt})'
           f'&"。改法：选中上一行的序号格（A 列）往下拖一格，再把上一行最右边的「核对」格（{chk_cols}）往下拖一格；或者在下面空行加、再排序"',
           fmt=INT, helpers=hs, sheet=SH_BASE)


# ───────────── 清单 ─────────────
def _list_header(ck, title_f, heads):
    ws = ck.ws
    r = ck.r
    ws.merge_cells(f'A{r}:{C["说明"]}{r}')
    put(ws, f'A{r}', title_f, F_SEC, fill(SEC_FILL), align=AL)
    ck.exempt(f'A{r}')
    for i in range(2, CI(C['说明']) + 1):
        ws.cell(row=r, column=i).border = BD
    ws.row_dimensions[r].height = 20
    header(ws, r + 1, list(zip(['A', 'B', 'C', 'D', 'E', 'F'], heads)), '核对', height=22)
    ck.r += 2


def _list(ck, title_txt, count_f, name, key, ftitle, ftext):
    """链式 MATCH 清单（前 CHK_LIST 笔）：隐藏 K 列＝在 _流 第几个位置；台帐第几行＝layout.z_row（v3.1：位置≠第几条）。
       标题：录入模式「先不列」；平时「共 x 笔」（x 放标题行的 K 格，pause0）"""
    ws = ck.ws
    K = CL(HELP0)
    tr = ck.r
    ws[f'{K}{tr}'] = pause0(_f(count_f) if count_f.startswith('=') else count_f)
    ws[f'{K}{tr}'].font = F_HELP
    _list_header(ck, f'=IF({ENTRY},"{title_txt}（录入模式中：先不列）","{title_txt}（共 "&${K}${tr}&" 笔，列前 {CHK_LIST} 笔）")',
                 ['台帐第几行', '日期', '账户（台帐里写的）', '收入', '支出', ftitle])
    r0 = ck.r
    for i in range(CHK_LIST):
        r = r0 + i
        k = f'${K}{r}'
        ws[f'{K}{r}'] = pause0(f'IFERROR(MATCH({key},{name},0),0)' if i == 0 else
                               f'IF({K}{r - 1}=0,0,{next_pos(f"{K}{r - 1}", key, name, N_FLOW)})')
        ws[f'{K}{r}'].font = F_HELP
        put(ws, f'A{r}', pause(f'IF({k}=0,"",{z_row(k)})'), F_TXT, FILL_NONE, fmt=INT, align=AC)
        put(ws, f'B{r}', pause(f'IF({k}=0,"",IF(INDEX(流_日期,{k})>0,INDEX(流_日期,{k}),""))'), F_TXT, FILL_NONE, fmt=DATE, align=AC)
        put(ws, f'C{r}', pause(f'IF({k}=0,"",{_tx("账户", k)})'), F_TXT, FILL_NONE, align=AL)
        put(ws, f'D{r}', pause(f'IF({k}=0,"",INDEX(流_收,{k}))'), F_TXT, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'E{r}', pause(f'IF({k}=0,"",INDEX(流_支,{k}))'), F_TXT, FILL_NONE, fmt=ACC, align=AR)
        put(ws, f'F{r}', pause(f'IF({k}=0,"",{ftext(k)})'), F_TXT, FILL_NONE, align=AL)
    ck.r = r0 + CHK_LIST + 1
    return r0


def _unk_list(ck):
    """五、认不出的写法（前 10 种）：先链式找前 UNK_N 笔「※ …「…」不在…」（流_问本），去重，再取前 10 种"""
    ws = ck.ws
    X, Y, Z, AA = (UNK_COL[k] for k in ('位置', '键', '首现', '累计'))
    b0, b1 = 5, 5 + UNK_N - 1
    for c, t in ((X, '认不出·位置'), (Y, '认不出·键'), (Z, '首现'), (AA, '累计种数')):
        ws[f'{c}4'] = t
        ws[f'{c}4'].font = F_HELP
    for i in range(UNK_N):
        r = b0 + i
        ws[f'{X}{r}'] = pause0(f'IFERROR(MATCH({UNK_PAT},流_问本,0),0)' if i == 0
                               else f'IF({X}{r - 1}=0,0,{next_pos(f"{X}{r - 1}", UNK_PAT, "流_问本", N_FLOW)})')
        ws[f'{Y}{r}'] = pause(f'IF({X}{r}=0,"",LEFT(INDEX(流_问本,{X}{r}),FIND("」",INDEX(流_问本,{X}{r}))))')
        # 首现＝这个写法第一次出现的那一行（MATCH 定长区域，不写往下撑的 $Y$5:Y{r−1}，WPS 存盘不拆）
        ws[f'{Z}{r}'] = pause0(f'IF({X}{r}=0,0,IF(IFERROR(MATCH({esc(f"{Y}{r}")},${Y}${b0}:${Y}${b1},0),0)=ROW()-{b0 - 1},1,0))')
        ws[f'{AA}{r}'] = pause0(f'{Z}{r}' if i == 0 else f'{AA}{r - 1}+{Z}{r}')
        for c in (X, Y, Z, AA):
            ws[f'{c}{r}'].font = F_HELP
    _list_header(ck, f'=IF({ENTRY},"五、认不出的写法（录入模式中：先不列）","五、认不出的写法（共 "&{CNT["认不出名字"]}&" 行；前 {UNK_N} 行里有 "'
                     f'&${AA}${b1}&" 种，列前 10 种）")',
                 ['序号', '种类', '台帐里的写法', '几行', '第一次在台帐第几行', '怎么改'])
    K, L_ = CL(HELP0), CL(HELP0 + 1)
    r0 = ck.r
    fix = {'账户': '去基础资料③：改过名就把这个写法加进那个账户的「别名」；是新卡就加一行',
           '收费项': '去基础资料④：把这个写法加进对的收费项的「别名」，或新加一个收费项；也可以在台帐改成下拉里的',
           '转账的另一户': '这一格只填本册的账户；外单位、个人写在收付人，收费项改成往来款、借款等',
           '归属公司': '填基础资料②里的公司简称'}
    fx = '""'
    for k_, t in reversed(list(fix.items())):
        fx = f'IF($B{{r}}="{k_}","{t}",{fx})'
    for i in range(10):
        r = r0 + i
        k, key = f'${K}{r}', f'${L_}{r}'
        ws[f'{K}{r}'] = pause0(f'IFERROR(MATCH({i + 1},${AA}${b0}:${AA}${b1},0),0)')
        ws[f'{L_}{r}'] = pause(f'IF({K}{r}=0,"",INDEX(${Y}${b0}:${Y}${b1},{K}{r}))')
        ws[f'{K}{r}'].font = ws[f'{L_}{r}'].font = F_HELP
        z = f'INDEX(${X}${b0}:${X}${b1},{k})'
        put(ws, f'A{r}', pause(f'IF({k}=0,"",{i + 1})'), F_TXT, FILL_NONE, align=AC)
        put(ws, f'B{r}', pause(f'IF({key}="","",MID({key},3,FIND("「",{key})-3))'), F_TXTB, FILL_NONE, align=AC)
        put(ws, f'C{r}', pause(f'IF({key}="","",MID({key},FIND("「",{key})+1,LEN({key})-FIND("「",{key})-1))'), F_TXT, FILL_NONE, align=AL)
        put(ws, f'D{r}', pause(f'IF({key}="","",IFERROR(COUNTIF(流_问本,{esc(key)}&"*"),0))'), F_TXT, FILL_NONE, fmt=INT, align=AC)
        put(ws, f'E{r}', pause(f'IF({k}=0,"",{z_row(z)})'), F_TXT, FILL_NONE, fmt=INT, align=AC)
        put(ws, f'F{r}', pause(f'IF({key}="","",{fx.format(r=r)})'), F_TXT, FILL_NONE, align=AL)
    ck.r = r0 + 11


# ═════════════════════════════ 数据整理记录 ═════════════════════════════
LOG_COLS = ['序号', '部分', '类别 / 项目', '原表位置 / 涉及', '字段', '原来', '现在 / 处理', '依据 / 说明']
LOG_W = dict(A=7, B=14, C=20, D=26, E=10, F=34, G=34, H=96)
PART = ['一、整理口径', '二、待确认', '三、账户对照', '四、模板这边的调整', '五、修改明细']
LOG_FILE = 'A074_数据整理记录.xlsx'


_ERR_TXT = re.compile(r'#(VALUE!|DIV/0!|REF!|NAME\?|NULL!|NUM!|N/A)')


def _num(s):
    """「20,187.05」→ 20187.05；认不出 None"""
    try:
        return float(str(s).replace(',', '').replace('，', ''))
    except (TypeError, ValueError):
        return None


def _m(s):
    v = _num(s)
    return f'{v:,.2f}' if v is not None else (str(s) if s not in (None, '') else '')


def _d(v):
    if isinstance(v, dt.date):
        return f'{v.year}/{v.month}/{v.day}'
    return '' if v in (None, '') else str(v)


def _pending_list_rows(p):
    """待确认里要逐笔列出来的（ctx.pending[i]['清单']）：放在那一条下面（复核 cls-02：47 笔对方不在本册，本册里就看得到是哪几笔）"""
    lst = p.get('清单') or []
    out = []
    for j, x in enumerate(lst, 1):
        inc, exp = _num(x.get('收入')) or 0, _num(x.get('支出')) or 0
        amt = '收 ' + _m(inc) if inc else ('支 ' + _m(exp) if exp else '')
        why = f'摘要「{x.get("摘要") or ""}」' + (f'；{x["说明"]}' if x.get('说明') else '')
        out.append((PART[1], f'└ 清单 {j}/{len(lst)}', x.get('原表位置', ''), _d(x.get('日期')),
                    f'{x.get("账户", "")}｜{amt}', f'记「{x.get("收费项", "")}」' if x.get('收费项') else '', why))
    return out


def _log_rows(ctx, detail=True):
    """数据整理记录各行 (部分, 类别, 位置, 字段, 原来, 现在, 依据)。detail=False（主本）：五 只写一行「在单独一本」"""
    rows = []
    for t, txt in ctx.get('clean_rules') or []:
        rows.append((PART[0], t, '', '', '', '', txt))
    for p in ctx.get('pending') or []:
        amt = p.get('涉及金额', '')
        rows.append((PART[1], p.get('问题', ''), p.get('涉及', ''), '涉及金额', _m(amt) if amt else '',
                     '已按「先这样处理」导进总表，等你确认', p.get('说明', '')))
        rows += _pending_list_rows(p)
    chk = {r.get('账户'): r for r in ctx.get('acct_check') or []}
    for a in ctx.get('accounts') or []:
        nm = a['账户名称']
        c = chk.get(nm)
        st = a.get('状态') or '在用'
        stop = a.get('停用日')
        extra = f'状态 {st}' + (f'，停用日 {stop.year}/{stop.month}/{stop.day}' if stop else '')
        if c:
            old = (f'期初 {_m(c.get("期初"))}｜原表 收 {_m(c.get("原表 收入"))}｜支 {_m(c.get("原表 支出"))}｜期末 {_m(c.get("原表 期末"))}'
                   f'（{c.get("原表笔数", "")} 笔）')
            same = all(_num(c.get(f'原表 {k}')) == _num(c.get(f'修正后 {k}')) for k in ('收入', '支出', '期末'))
            if not same:
                old += f'；修正后 收 {_m(c.get("修正后 收入"))}｜支 {_m(c.get("修正后 支出"))}｜期末 {_m(c.get("修正后 期末"))}'
            new = (f'收 {_m(c.get("新 收入"))}｜支 {_m(c.get("新 支出"))}｜期末 {_m(c.get("新 期末"))}（{c.get("新行数", "")} 行）')
            d = _num(c.get('新−修正后 期末'))
            why = (f'新总表期末－修正后期末＝{_m(c.get("新−修正后 期末"))}' + ('（对上了）' if d == 0 else '（对不上，见待确认）') +
                   f'；{extra}' + ('；' + a['备注'] if a.get('备注') else ''))
        else:
            old, new = f'期初 {_m(a.get("期初余额"))}', ''
            why = extra + ('；' + a['备注'] if a.get('备注') else '')
        rows.append((PART[2], nm, a.get('原表名') or (a.get('别名') or '').split('、')[0], '期初/收/支/期末', old, new, why))
    for it in ctx.get('items') or []:
        m = re.search(r'属性原为「(.+?)」，模板改为「(.+?)」', it.get('说明') or '')
        if m:
            rows.append((PART[3], '收费项属性', it['收费项'], '属性', m.group(1), m.group(2),
                         '项目建设算「投资建设」、待定算「待分类」，都单列、不算经营收支（这样分对不对，要你确认）'))
    items = ctx.get('items') or []
    if items:
        rows.append((PART[3], '收费项顺序', '基础资料④', '顺序', '导入时的顺序（原付费项的顺序）',
                     '收入 → 内部转账 → 支出 → 投资建设 → 往来 → 待分类，同属性按笔数从多到少', '下拉里常用的在上面（房租、物业、水电收入排最前）'))
    op = next((it for it in items if it['收费项'] == ITEM_OPEN), None)
    if op is not None and not op.get('_笔数'):
        rows.append((PART[3], '预置收费项', ITEM_OPEN, '收费项', '（没有）', f'{ITEM_OPEN}（往来，大类「{op.get("大类", "")}」）',
                     '年中新开的账户期初填 0，开户那天的余额在台帐录一笔「账户转入」，不改历史月份的资金总额'))
    if ctx.get('parties'):
        rows.append((PART[3], '收付人顺序', '基础资料⑤', '顺序', '导入时的顺序', '按笔数从多到少', '下拉里常用的在上面'))
    for a in ctx.get('accounts') or []:
        if a.get('状态') in ('暂停', '销户') and a.get('停用日'):
            d = a['停用日']
            rows.append((PART[3], '停用日', a['账户名称'], '停用日', '（原表没有）', f'{d.year}/{d.month}/{d.day}',
                         f'状态「{a["状态"]}」：停用日＝最后一笔流水的日期，之前的流水不提示「暂停户还在记」'))
    log = ctx.get('cleanlog') or []
    if detail:
        for x in log:
            rows.append((PART[4], x.get('类别', ''), x.get('原表位置', ''), x.get('字段', ''), x.get('原值', ''), x.get('新值', ''),
                         x.get('依据', '')))
    elif log:
        fn = ctx.get('log_file') or LOG_FILE
        rows.append((PART[4], '在单独一本', f'《{fn}》', '', f'{len(log):,} 条', '',
                     f'修改明细 {len(log):,} 条（原表哪一格、原来是什么、改成什么、依据）放在同一文件夹的《{fn}》里，'
                     f'那一本还有上面一～四的全部内容。日常记账不用看，核对导入的数时打开（点「部分」「类别」的筛选箭头只看某一类）'))
    return rows


def _lines(v, width):
    """文字在这么宽的格子里大约占几行（中文一个字约 2 个宽度单位）"""
    if v is None or v == '':
        return 1
    per = max(4, int(width / 2.05))
    return sum(max(1, -(-len(seg) // per)) for seg in str(v).split('\n'))


def _import_date(ctx):
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), '导入', 'out', 'entries.json')
    try:
        return dt.date.fromtimestamp(os.path.getmtime(p))
    except OSError:
        return None


def build_log(ws, ctx, detail=False):
    """数据整理记录（静态，全是值）。detail=False：主本（一～四＋「五在单独一本」一行）；detail=True：单独一本（五部分全有，build.py 写）。
       返回表里的行数"""
    widths(ws, LOG_W)
    es = [e for e in ctx.get('entries') or [] if e]               # 场景里的空行（None）不算
    ds = [e['日期'] for e in es if isinstance(e.get('日期'), dt.date)]
    fn = ctx.get('log_file') or LOG_FILE
    n = dict(规则=len(ctx.get('clean_rules') or []), 待确认=len(ctx.get('pending') or []), 账户=len(ctx.get('accounts') or []),
             明细=len(ctx.get('cleanlog') or []))
    if detail:
        tip = ('这一本是《A074 集团多账户资金台帐》导入原来 2026 年日记账时的完整记录：这本账改了什么、为什么改、还有哪些要你确认。'
               '一张表分五段（看「部分」列）：一、整理口径（总的做法）；二、待确认（先按一种办法导进来了，要你拍板；要逐笔看的列在那一条下面）；'
               '三、账户对照（每个账户原表和新总表的收、支、期末）；四、模板这边的调整；五、修改明细（逐格：原表哪一格、原来是什么、改成什么、依据）。'
               '第 4 行有筛选：点「部分」「类别」的箭头只看某一段、某一类。这一本是导入时写死的，不会跟着台帐变。')
    else:
        tip = ('导入原来 2026 年日记账时，这本账改了什么、为什么改、还有哪些要你确认（要点）。分几段（看「部分」列）：一、整理口径（总的做法）；'
               '二、待确认（先按一种办法导进来了，要你拍板；要逐笔看的列在那一条下面，按「原表位置」在【资金台帐】P 列能筛到）；'
               '三、账户对照（每个账户原表和新总表的收、支、期末）；四、模板这边的调整；'
               f'五、逐格的修改明细在同一文件夹的《{fn}》里（日常不用看）。第 4 行有筛选。'
               '这张表是导入时写死的，不会跟着台帐变；开新本时清空或写「上一年的导入记录见上一本」。')
    title(ws, '数据整理记录', 'H', tip, h2=44)
    imp = _import_date(ctx)
    if n['明细'] or n['待确认']:
        info = ('导入来源：整理后的原 2026 年日记账' + (f'；导入日期 {imp.year}/{imp.month}/{imp.day}' if imp else '') + '；' +
                (f'流水 {len(es):,} 笔（{min(ds).year}/{min(ds).month}/{min(ds).day}～{max(ds).year}/{max(ds).month}/{max(ds).day}）；'
                 if ds else '') +
                f'整理口径 {n["规则"]} 条、待确认 {n["待确认"]} 条、账户 {n["账户"]} 个、修改明细 {n["明细"]:,} 条' +
                ('' if detail or not n['明细'] else f'（在《{fn}》）'))
    else:
        info = '（这本没有导入记录：空白模板，或者开新本时清空了）'
    if detail:
        put(ws, 'A3', None, F_KPI_L, FILL_KPI, align=AL)
    else:
        home_link(ws, 'A3')
    ws.merge_cells('B3:H3')
    put(ws, 'B3', info, F_KPI_L, FILL_KPI, align=AL)
    ws.row_dimensions[3].height = 24
    hr = 4
    header(ws, hr, list(zip('ABCDEFGH', LOG_COLS)), '核对')
    rows = _log_rows(ctx, detail=detail)
    sec_fill = {PART[0]: fill('FFE2EFDA'), PART[1]: fill('FFFFF2CC'), PART[2]: fill('FFDDEBF7'), PART[3]: fill('FFEDE2F6'),
                PART[4]: FILL_NONE}
    al_c = Alignment(horizontal='center', vertical='top', wrap_text=True)
    for i, row in enumerate(rows, 1):
        r = hr + i
        vals = (i,) + row
        part = row[0]
        big = part != PART[4] or not detail
        sub = isinstance(row[1], str) and row[1].startswith('└')
        for j, v in enumerate(vals):
            if isinstance(v, str):
                v = _ERR_TXT.sub(lambda m_: '＃' + m_.group(0)[1:], v)   # 原表里抄来的「#REF!」是文字：换全角＃，免得被当成公式错误数进去
            c = ws.cell(row=r, column=j + 1, value=v if v != '' else None)
            if isinstance(v, str) and v.startswith('='):
                c.data_type = 's'                  # 原值、依据里以 = 开头的是文字（原表的公式），不当公式
            c.font = F_TXTB if (j == 2 and big and not sub) else F_TXT
            c.border = BD
            c.alignment = al_c if j in (0, 1, 4) else WRAP_TOP
            if part in sec_fill and (part != PART[4] or not detail):
                c.fill = sec_fill[part]
        if big:
            ws.row_dimensions[r].height = min(409, 14 * max(_lines(v, LOG_W[CL(j + 1)]) for j, v in enumerate(vals)) + 6)
    last = hr + max(1, len(rows))
    ws.auto_filter.ref = f'A{hr}:H{last}'
    ws.freeze_panes = f'A{hr + 1}'
    print_setup(ws, f'{hr}:{hr}')
    return len(rows)


# ═════════════════════════════ 使用说明 ═════════════════════════════
def _help_text(ctx):
    n_acc = len(ctx.get('accounts') or [])
    n_co = len(ctx.get('companies') or [])
    n_e = len([e for e in ctx.get('entries') or [] if e])
    n_rl = len(ctx.get('rules') or [])
    n_pd = len(ctx.get('pending') or [])
    n_log = len(ctx.get('cleanlog') or [])
    fn = ctx.get('log_file') or LOG_FILE
    lc, ls = LEDGER_COLS, LEDGER_SHOW
    in_cols = '、'.join(f'{a}～{b}' if a != b else a for a, b in _ledger_input_spans())
    show = f'{ls["余额"]}、{ls["问题"]}～{ls["收付人认"]}'
    spare = '」「'.join(s for s, _c in SPARE_SHEETS)
    S = []

    def sec(t, items):
        S.append((t, items))
    sec('一、这本账有什么', [
        ('一张总表', f'所有银行卡、现金账户的每一笔收付都记在【资金台帐】一张表里（现在 {n_acc} 个账户、{n_co} 家公司'
                  + (f'，已导入 {n_e:,} 笔' if n_e else '') + '）。其余的表全是从它自动算出来的，不用填。'),
        ('表的顺序', '首页（老板看资金）→ 资金台帐（出纳每天录）→ 基础资料（账户、公司、收费项、收付人、摘要规则）→ 8 张报表 → '
                 f'账户查询、每个账户一张子表、备用子表 {len(SPARE_SHEETS)} 张 → 核对表、数据整理记录 → 使用说明。'
                 + (f'导入时逐格的修改明细（{n_log:,} 条）在同一文件夹的《{fn}》里，平时不用打开。' if n_log else '')),
        ('颜色', '淡黄色的格子要填；黄底红字的是报表上的选择格（年份、日期、月份，可以改）；灰色的是公式，别动。'
               '※ 红底＝这一行有错、必须改；△ 琥珀色字＝照算了、最好补；琥珀色斜体的收费项、收付人＝按摘要猜出来的。'),
        ('两种口径', '账户口径（子表、账户查询、资金日报、各账户余额）：收入、支出就是银行流水上进出的钱，只是把内部转账单列。'
                 '经营口径（首页收支、月度汇报表、收支分类、收支报表……）：只算收费项属性是「收入」「支出」的；'
                 '投资建设（建项目、买设备）、往来（借款、还款、押金、投资款）、待分类（待定的、认不出的）、内部转账都单列，不算经营收支。'),
    ])
    sec('二、每天怎么录（资金台帐）', [
        ('一笔一行', '从最下面第一个空行接着录。日期、账户必填（淡黄色两列）。'),
        ('日期', '写成 2026/10/13 或 2026-10-13 这样的真日期。写成 10.13、10-13（没有年）的认不出，会 ※。'),
        ('账户', '从下拉选；写原来表名（比如「易佰兴私户」）、旧名也认得（基础资料③「别名」里有的都认）。'),
        ('收入、支出', '只填一个。银行扣的手续费另起一行，收费项选「银行手续费」。金额填负数会自动当另一栏算（△ 提醒一下）。'),
        ('收费项', '最好从下拉选。空着也行：会按基础资料⑥的摘要关键词去猜，猜出来的在 N 列显示琥珀色斜体，猜不出来的算「待分类」。'),
        ('一笔钱拆几行', '一笔银行流水里有几样钱（比如 C118 交 1 月租金 1,500＋物业 150＋水电 105.6＝1,755.6）：录 3 行。'
                    '第 1 行录好以后，在下一行（新的那一行）选中 A～E 按 Ctrl+D，就把上一行抄下来了，再改收费项和金额；第 3 行同样。'
                    '注意：只选一行按 Ctrl+D 是「把上一行抄进选中的这一行」，所以要在新的那一行按，别选已经录好的那一行。'
                    '要夹在中间录：先点左边行号选整行，右键「插入」2 行，再这样抄。银行流水上的 1,755.6 对的是这 3 行合计。'
                    '摘要一样、收费项空着的几行猜不准，一定要选收费项。'),
        ('收付人', '对方是谁（租户、客户、供应商、个人）。空着会按摘要猜；⑤里没有的照记，报表里算「没填」。'),
        ('归属公司', '只有私户替别的公司收付时才填（比如个人卡替力喆实业收的租金）：资金还算账户所属公司，收支算到归属公司。'),
        ('备注', '开票没开票、哪个项目、经办人……都写在备注。以后在备注列点筛选「包含 未开票」就能查出来。'),
        ('录完看一眼', f'{ls["余额"]} 列是这个账户录到这一行的余额，跟网银余额对一下；{ls["问题"]} 列是这一行的问题：※ 必须改，△ 看一眼。'),
        ('录很多时', '一次要录几十上百笔，先去【基础资料】①把「录入模式」填「是」，录完清空（见「十二」）。'),
    ])
    sec('三、内部转账（集团里两个账户之间互转）', [
        ('记两行', '转出的账户记一笔支出、转入的账户记一笔收入，收费项都选「内部转账」，「转账的另一户」（'
                f'{lc["对方账户"]} 列）填对方账户。收费项空着、只填了另一户，也自动当内部转账。'),
        ('允许差几天', f'两头的日期可以差 {PAIR_DAYS} 天以内（跨行转账第二天才到），金额要一样。只录了一头、金额或日期对不上的，'
                  f'{ls["问题"]} 列 △「内部转账没配上」，核对表「六」有清单。'),
        ('对方不在本册', '对方是外单位、个人、或者本册没有的卡：不算内部转账，收费项选借款转入、还款、往来款这类，对方名字写在收付人，'
                   f'{lc["对方账户"]} 列别填（填了会 ※）。'),
        ('报表里', '内部转账不算收入支出，各表单列「内部转入」「内部转出」；全集团一年下来净额应该是 0（核对表查）。'),
    ])
    sec('四、网银流水怎么粘', [
        ('先排好列', '在网银导出的表里先把列排成台帐的顺序：日期、账户、收费项、收付人、摘要、收入、支出；右边是转账的另一户、归属公司、备注。'),
        ('选择性粘贴', f'复制以后，在台帐第一个空行右键「选择性粘贴 → 数值」，只粘到 {in_cols} 列。'
                  f'{show} 是自动列，不要粘上去（盖掉了核对表会报）；右边 {LEDGER_ANCHOR} 列（勿删）也别粘。'),
        ('文字也认', '网银里的文字日期（2026-10-01 14:23:05、2026.10.1、2026年10月1日）、带逗号或 ¥ 的金额、「-」都认得，'
                 f'{ls["问题"]} 列 △ 提醒是文字，最好改成真日期、数字。'),
        ('一栏金额', '只有一栏「交易金额」、带正负号的（-5000），整列粘到收入栏：负数自动当支出算。'),
        ('账户一列', '整批都是同一个账户：在第一行选账户，再把第一行连同下面要填的格子一起选中，按 Ctrl+D 填满。'),
        ('粘完看重复', f'同账户、同一天、同金额、同摘要的两笔（前后 {DUP_WIN} 行内）会 △「重复粘贴了？」。真的是两笔，就在备注写一句。'),
    ])
    sec('五、收费项按摘要认、固化', [
        ('怎么认', '收费项空着时，按基础资料⑥「摘要关键词规则」从上往下找，摘要（或收付人）里有关键词、方向和适用范围也对上的，'
                '第一条就算；猜出来的在台帐 N 列显示琥珀色斜体，子表里也是。'
                + (f'现在有 {n_rl} 条规则。' if n_rl else '')),
        ('猜错了', '直接在收费项（C 列）选对的；常出现的新摘要，在⑥加一条规则（越具体的放越上面）。'
                '同一个摘要能对上好几种收费项（一行拆成租金、物业、水电时摘要一样）会 △，请直接选。'),
        ('固化', f'月底核对完、※ 都改完以后：先取消所有筛选；选中 {ls["收费项认"]}5 到最后一笔流水那一行'
               '（别点列标选整列，也别拉到下面的空行），复制，点 C5，右键「选择性粘贴 → 数值」；'
               f'{ls["收付人认"]} 列同样粘到 D5。{ls["收费项认"]}、{ls["收付人认"]} 两列认不出的照抄 C、D 列原来的字，粘回去不会把名字抹掉。'
               '别先筛选再粘：往筛选过的区域粘，会连着往下写，把中间藏起来的行也改掉。'
               '固化以后再改规则，不会悄悄改掉已经结了的月份；空着的行越少，改一格越快。'),
    ])
    sec('六、看哪几张表', [
        ('首页', '老板看：顶上一排提醒（※、△、内部没配上、待分类、几户几天没记账…点一下跳过去）；资金总额（在用 / 暂停）；'
               '截止那天、本月、本年的经营收支；各公司资金；12 个月走势；每个账户的余额、最后一笔是哪天。黄格「截止日期」可以改。'),
        ('资金日报', '① 一段日期（最多 31 天）每个账户每天收付；② 某一天各账户：昨日余额、本日收支、本日余额；③ 某个月各账户。'),
        ('月度汇报表', '一年 12 个月的经营收入、支出、投资建设、往来、内部、月末资金；下面选一个月看各账户。'),
        ('公司汇总', '每家公司一年的年初资金、经营收支、年末资金；公司 × 月；公司 → 账户余额。'),
        ('收支分类报表', '按属性 → 大类 → 收费项展开每个月；黄格可以选一家公司（按归属公司）。'),
        ('收支分项汇总', '收费项 × 公司（像原来日记账的【汇总】）：选月份或看全年。'),
        ('收支报表', '收入按收付人、支出按收费项和收付人排名，按收付人类型合计。'),
        ('任意时间段', '任意两天之间的经营收支、各账户期初期末、按收费项和收付人排名。'),
        ('年末余额', '某一年年末每个账户的余额；下面是开新本要复制的那一块。'),
        ('子表、账户查询', '每个账户一张子表（像原来的日记账：月、日、收费项、摘要、收入、支出、余额），右边是 12 个月的小表；'
                    '账户查询可以选任意账户、某个月。子表是自动分出来的，别在子表里改，要改去台帐改。'),
        ('核对表', '整本账算得对不对、台帐里还有什么要补、插删行有没有把台帐弄乱；下面有清单（台帐第几行）。'),
        ('数据整理记录', (f'导入原日记账时改了什么、依据、待你确认的 {n_pd} 条（要点）；逐格的修改明细在《{fn}》。' if n_pd
                     else '导入时改了什么（空白模板没有）。')),
    ])
    sec('七、年份怎么查', [
        ('报表的年份格', '每张报表、子表顶上都有黄色的「年份」：空着＝管理年度（基础资料①；那里也空着＝最后一笔流水的年份）。'
                    '填 2026、26 或那年的某一天都认。'),
        ('1 月看去年', '1 月份还想看去年全年：把基础资料①「管理年度」填上一年，各报表、子表、首页都跟着看那一年（首页截止到那年 12/31）。'),
        ('首页截止日期', '首页黄格填任意一天，资金总额、各账户余额都算到那天；空着＝最后一笔流水那天。'),
        ('一年一本', '一本只记一年（开账日那年，再加次年 1 月还没开新本时先记在这里的）。看上一年的报表打开上一本。'),
    ])
    sec('八、开新本（每年 1 月，按顺序做）', [
        ('①', '把这一本另存一份，文件名改成新的一年（比如「…2027.xlsx」）。下面都在新文件里做。'),
        ('②', '【年末余额】黄格年份填旧的那一年（比如 2026）→ 选中下面「粘贴区」整块复制 → 到【基础资料】③账户「账户名称」列第一个账户那格，'
              '右键「选择性粘贴 → 数值」（各账户的期初余额就是去年年末余额）。'),
        ('③', '【基础资料】①「开账日」改成新一年的 1 月 1 日，「管理年度」清空。'),
        ('④', f'【资金台帐】点日期列的筛选箭头 → 只勾旧的那一年 → 选中筛出来的这些行的 {in_cols} 列按 Delete 清掉内容'
              f'（别删行；别动 {show} 公式列和 {LEDGER_ANCHOR} 列）→ 取消筛选，再按日期升序排一次（新一年的行就到上面了）。'),
        ('⑤', '【数据整理记录】清空，或写一句「上一年的导入记录见上一本」。核对表「开新本后旧年流水没删」应该是 ✓。'),
    ])
    sec('九、加账户、收费项、收付人、规则', [
        ('往下加', '都在基础资料各块下面的空行接着加；③账户可以在中间插行、可以排序。'),
        ('中间插了行', '基础资料中间插出来的行没有序号和核对公式：选中上一行的序号格（A 列）往下拖一格，再把上一行最右边的「核对」格往下拖一格；'
                  '或者干脆在下面空行加、再排序。忘了的话核对表「基础资料中间插的行没带公式」会报。'),
        ('改名', '账户、收费项、收付人改名时，把旧名字加进那一行的「别名」（多个用「、」隔开）：台帐里以前的写法照样认得。'),
        ('年中新开的账户', '期初余额填 0，开户那天的余额在台帐录一笔收费项「账户转入」。这样不会改到以前月份的资金总额。'),
        ('不用的账户', '状态改「销户」（或「暂停」），填停用日；别删行（删了以后那个账户的流水全变 ※）。'),
        ('子表', '每张子表靠③的「子表名」认账户：子表名那一列别改，子表别删、别改名（首页、核对表的链接按表名跳；不用的可以隐藏）。'
               f'新开的卡先用「{spare}」：在 D3 选账户就行，表名别改，用的是哪个账户看它第 1 行的标题。'
               '备用的也用完了：复制任意一张子表，在新表 D3 选账户（复制出来的表核对表、首页不认，所以先用备用的）。'),
        ('首页藏起来的行', '首页、报表里按现在的公司、账户数把多余的行藏起来了：加了公司、账户以后，选中那几行右键「取消隐藏」。'),
    ])
    sec('十、插行、删行、排序（一定照这样做）', [
        ('点行号选整行', '插行、删行一律先点左边的行号选中整行，再右键「插入」「删除」。别只选几个格子插删'
                    '（「插入 → 活动单元格下移」「删除 → 下方单元格上移」都不行），也别只剪切 A～P 的格子：台帐会被挪乱，首页和台帐第 1 行会红字报。'),
        ('插了行补公式', f'中间插了行要在新行里录：先选中上一行的 {ls["余额"]}、{ls["问题"]}、{ls["属性"]}、{ls["收费项认"]}、{ls["收付人认"]} 五格往下拖一格'
                    '（把公式带下来）；忘了的话核对表「显示列缺公式」会报。平时在最下面接着录，录完再排序，就不用补。'),
        ('挪行用排序', '要调顺序，点「日期」列表头的筛选箭头 → 升序（别选中整张表排序，会把第 1～3 行也排进去）。'
                  f'{ls["余额"]} 列按表里从上到下算，排好序就对了；报表、子表的月份表都按日期算，不受顺序影响。要剪切就点行号剪整行。'),
        ('空行别删', '不要的行选中按 Delete 清内容就行；一次删一大片（100 行以上）也先清内容。别「筛选空白 → 删除行」把下面的空行删光：'
                  '能录的行就少了（台帐第 3 行「还能录约 x 行」会变小）。'),
        ('别用这几样', '别用「数据 → 删除重复项」，别把台帐转成「表格 / 超级表」（Ctrl+T）：它们只挪格子不挪整行，跟只插删格子一样会把台帐弄乱。'
                   '怕重复粘贴看 L 列 △「重复粘贴了？」，手动删整行。'),
        (f'{LEDGER_ANCHOR} 列（勿删）', f'台帐最右边窄窄的 {LEDGER_ANCHOR} 列是公式定位用的：别填、别删、别在它上面粘东西。'),
        ('删行', '不要的流水点行号删整行，删了以后不用补任何东西；一次删 100 行以上先选中按 Delete 清内容（见上面「空行别删」）。'),
    ])
    sec('十一、首页、台帐第 1 行的红字（插删行报警）', [
        ('「表头或 Q 列被删了」', f'台帐第 4 行表头或右边 {LEDGER_ANCHOR} 列被删了，所有数都没算：马上撤销（Ctrl+Z）；已经存盘了找我们。'),
        ('「被插/删了单元格」', '只插删了格子（不是整行）、只剪了 A～P 的格子、删了中间的列、用了删除重复项：后面的数全乱了。马上撤销（Ctrl+Z）；'
                       '已经存盘了，按原样把那几格挪回去，或者找我们。核对表「台帐分段」写了大概是第几段（台帐第几行）。'),
        ('「有 x 行没算进来」', f'一段里一次插了 {SEG_GROW_MAX} 行以上、删了整段又在接缝插行、或者录到了容量以外：核对表「台帐分段」写了是台帐第几行。'
                       '改法：把这几行复制 → 粘到台帐最下面的空行 → 点行号删掉原来那几整行 → 按日期排序。'),
        ('「还能录约 x 行」「台帐满了」', '年底按「八」开新本；年中就满了找我们放大。'),
        ('「【录入模式】…」', '录完去基础资料①把「录入模式」清空。'),
    ])
    sec('十二、录很多时、打开慢、保存慢', [
        ('录入模式', '一次要录几十上百笔时，把基础资料①「录入模式」填「是」：子表、报表、核对表、首页先不算（这些表第 1 行变成红字「录入模式中」），'
                 '台帐的余额、※△ 照算，改一格快很多。录完清空（清空那一下要整本算几秒），报表就都回来了。'
                 '忘了清空，老板看到的报表是空的，所以录完一定清空。'),
        ('别用手动计算', '不要把 Excel / WPS 改成「手动计算」：这个设置会跟着文件存下来，别人打开或改首页截止日期时数不动、看到的是旧数。'
                   '要快就用上面的「录入模式」。'),
        ('保存慢', 'WPS：「文件 → 备份与恢复 → 备份中心 → 本地备份设置」改成「增量备份」，或者把定时备份的间隔拉长（每次定时备份都要把整本写一遍）。'
                'Excel：「文件 → 选项 → 保存」里「保存自动恢复信息时间间隔」调到 20～30 分钟。'),
        ('打开慢', 'Excel 2016 及更早的版本、WPS 打开时会先把整本算一遍，要等一会儿；关的时候问「是否保存」点保存，下次打开就快。'
                '手机、微信里预览看到的是最后一次保存时的数。'),
    ])
    sec('十三、核对表怎么看', [
        ('判断', '✓ 没问题；△ 看一眼、最好补；✗ 要改。第 3 行是总数。录入模式时只留第 3 行台帐的 ※、△ 行数，别的先不算。'),
        ('勾稽 ✗', '报表算的跟台帐直接算的对不上：多半是有人改了报表里的公式（报表、子表都锁着，只能改黄格），请告诉我们。'),
        ('台帐分段', f'台帐每 {SEG_ROWS} 行一段读进隐藏的计算表。表头或 {LEDGER_ANCHOR} 列被删、插删了格子（不是整行）、一段一次插得太多、'
                 '录到了容量以外，这一项会 ✗，说明里写了是台帐第几行、怎么改（同「十一」）；整段删掉了（一次删了 100 多行）是 △：数没错，只是能录的行少了。'),
        ('待办', '台帐里要补要改的（※、△、按摘要认的、待分类、没配上……）；下面清单写了是台帐第几行（在台帐按 Ctrl+G，输入 A 和行号就跳过去）。'
               '待分类没有清单：在台帐 M 列（属性）点筛选只勾「待分类」。'),
        ('子表', '每张子表一行：余额、笔数、明细跟台帐比；「放不下」的看【账户查询】选月份。'),
    ])
    return S


def _ledger_input_spans():
    """资金台帐录入列连成的几段（显示列不算）：[(起列, 止列)]"""
    cols = sorted(CI(c) for c in LEDGER_COLS.values())
    spans, s = [], cols[0]
    for a, b in zip(cols, cols[1:] + [None]):
        if b != a + 1:
            spans.append((CL(s), CL(a)))
            s = b
    return spans


def build_help(ws, ctx):
    widths(ws, dict(A=4, B=16, C=110))
    tip = '先看「一」「二」；出纳每天用到的是「二」～「五」和「十」，老板看「六」「七」，每年 1 月看「八」；台帐、首页出了红字看「十一」。'
    title(ws, '使用说明', 'C', tip, h2=24)
    home_link(ws, 'A3')
    put(ws, 'B3', '目录', F_KPI_L, FILL_KPI, align=AC)
    S = _help_text(ctx)
    put(ws, 'C3', '　'.join(t for t, _i in S), F_NOTE, FILL_NONE, align=ALW)
    ws.row_dimensions[3].height = 44
    r = 5
    hf = fill('FF1F3864')
    for t, items in S:
        ws.merge_cells(f'A{r}:C{r}')
        put(ws, f'A{r}', t, F_SEC, hf, align=AL)
        ws.row_dimensions[r].height = 22
        r += 1
        for i, (k, v) in enumerate(items, 1):
            put(ws, f'A{r}', i, F_AUTO, FILL_NONE, align=Alignment(horizontal='center', vertical='top'))
            put(ws, f'B{r}', k, F_TXTB, FILL_NONE, align=WRAP_TOP)
            put(ws, f'C{r}', v, F_TXT, FILL_NONE, align=WRAP_TOP)
            ws.row_dimensions[r].height = max(18, 14 * max(_lines(v, 110), _lines(k, 16)) + 6)
            r += 1
        r += 1
    put(ws, f'A{r}', None, border=False)
    ws.merge_cells(f'A{r}:C{r}')
    put(ws, f'A{r}', '有不明白的、数对不上的，把核对表截图发给我们。', F_NOTE, FILL_NONE, align=AL, border=False)
    ws.freeze_panes = 'A5'
    print_setup(ws, None, landscape=False)
    return r


def build(wb, ctx):
    build_chk(wb[SH_CHK], ctx)
    build_log(wb[SH_LOG], ctx, detail=False)
    build_help(wb[SH_HELP], ctx)
