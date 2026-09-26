# -*- coding: utf-8 -*-
"""【资金台帐】收支混合录入 + 即时余额；【收付款明细】改成从资金台帐自动提取。"""
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from copy import copy
from openpyxl.styles import Font, PatternFill

from common import *

F_IN_BG  = fill('FFF3FAEF')   # 收入列淡绿
F_OUT_BG = fill('FFFFF8E1')   # 支出列淡黄


def _cat_lookup(col_name_cell, what_col):
    return (f'IFERROR(INDEX({SH_BASE}!${what_col}${CAT_R0}:${what_col}${CAT_R1},'
            f'MATCH({col_name_cell},{SH_BASE}!${C_NAME}${CAT_R0}:${C_NAME}${CAT_R1},0)),"")')


def build_cash(wb, ctx):
    ws = wb.create_sheet(SH_CASH)
    ws.sheet_properties.tabColor = C_CASH[2:]
    last = K_LAST
    title(ws, '资金台帐（收支混合录入 · 即时余额）', last, C_CASH,
          '💡 一笔一行，按发生顺序往下记：收钱填【收入金额】，付钱填【支出金额】，右边【账户余额】【总余额】自动滚出来。'
          '收客户的钱选【客户】、付厂家货款选【供应商】、花钱选【费用项目】（报销的填上【报销人】）——总览汇总、收付款明细、'
          '全量费用明细总表、收款码到账对账、对账单、利润表都从这里自动取数。选了客户/供应商/费用项目的，【收支项目】可以不选，会自动认。')

    # ── 第 3 行：汇总条（照截图：月份 / 本月收入 / 本月支出 / 上月结余 / 本月结余 / 当前总余额）──
    ws.row_dimensions[3].height = 30
    for i in range(1, CI(last) + 1):
        put(ws, f'{CL(i)}3', None, fill_=FILL_KPI, border=False)
    rng = lambda col: f'{col}${CASH_R0}:${col}${CASH_R1}'
    acc0 = f'SUM({SH_BASE}!${A_BAL0}${ACC_R0}:${A_BAL0}${ACC_R1})'
    kpis = [
        ('A', '月份：', 'B', ctx['kpi_month'], '0"月"', True),
        ('C', '本月收入：', 'D',
         f'=SUMIFS(${rng(K_IN)},${rng(K_DATE)},">="&DATE(年度,$B$3,1),${rng(K_DATE)},"<"&DATE(年度,$B$3+1,1),${rng(K_TO)},"<>{TO_XFER}")', MONEY2, False),
        ('E', '本月支出：', 'F',
         f'=SUMIFS(${rng(K_OUT)},${rng(K_DATE)},">="&DATE(年度,$B$3,1),${rng(K_DATE)},"<"&DATE(年度,$B$3+1,1),${rng(K_TO)},"<>{TO_XFER}")', MONEY2, False),
        ('G', '上月结余：', 'H',
         f'={acc0}+SUMIFS(${rng(K_IN)},${rng(K_DATE)},">="&年初日,${rng(K_DATE)},"<"&DATE(年度,$B$3,1))'
         f'-SUMIFS(${rng(K_OUT)},${rng(K_DATE)},">="&年初日,${rng(K_DATE)},"<"&DATE(年度,$B$3,1))', MONEY2, False),
        ('I', '本月结余：', 'J',
         f'={acc0}+SUMIFS(${rng(K_IN)},${rng(K_DATE)},">="&年初日,${rng(K_DATE)},"<"&DATE(年度,$B$3+1,1))'
         f'-SUMIFS(${rng(K_OUT)},${rng(K_DATE)},">="&年初日,${rng(K_DATE)},"<"&DATE(年度,$B$3+1,1))', MONEY2, False),
        ('K', '当前总余额：', 'L', f'={acc0}+SUM(${rng(K_IN)})-SUM(${rng(K_OUT)})', MONEY2, False),
    ]
    for lc, lab, vc, val, fmt, is_in in kpis:
        put(ws, f'{lc}3', lab, F_KPI_L, FILL_KPI, align=AR_, border=False)
        put(ws, f'{vc}3', val, Font(name=YH, sz=11, bold=True, color='FFC00000') if not is_in else Font(name=YH, sz=13, bold=True, color='FF1F4E79'),
            FILL_SEL if is_in else FILL_KPI, fmt, AL if not is_in else AC, border=is_in)
    put(ws, 'M3', '按日期统计，本月收入/支出不含内部转账', F_NOTE, FILL_KPI, align=ALW, border=False)
    ws.merge_cells(f'M3:{last}3')
    add_list_dv(ws, 'B3', '"1,2,3,4,5,6,7,8,9,10,11,12"', '选月份，看这个月的收支')

    # ── 第 4 行：蓝带 ──
    ws.row_dimensions[4].height = 6
    for i in range(1, CI(last) + 1):
        put(ws, f'{CL(i)}4', None, fill_=FILL_BAND, border=False)

    # ── 第 5 行：分组色带（截图里的「＋ 收入明细」绿、「＋ 支出明细」黄）──
    ws.row_dimensions[5].height = 30
    bands = [('A', 'E', '登 记 信 息', 'FF5B9BD5'), ('F', 'F', '＋ 收 入', C_INCOME), ('G', 'G', '＋ 支 出', C_EXPENSE),
             ('H', 'I', '即时余额（自动）', 'FF5B9BD5'), (K_CUS, K_WHO, '往来 / 费用（筛选列）', 'FF7030A0'),
             (K_NOTE, K_CHK, '备注 / 校验', 'FF5B9BD5')]
    for c0, c1, t, col in bands:
        if c0 != c1:
            ws.merge_cells(f'{c0}5:{c1}5')
        put(ws, f'{c0}5', t, Font(name=YH, sz=13, bold=True, color='FFFFFFFF'), fill(col), align=AC)
        if c0 != c1:
            for cc in range(CI(c0) + 1, CI(c1) + 1):
                put(ws, f'{CL(cc)}5', None, fill_=fill(col))

    # ── 第 6 行：表头 ──
    hdr = [(K_SEQ, '序号'), (K_DATE, '日期'), (K_ACC, '账户'), (K_CAT, '收支项目'), (K_MEMO, '摘要'),
           (K_IN, '收入金额'), (K_OUT, '支出金额'), (K_ABAL, '账户余额'), (K_TBAL, '总余额'),
           (K_CUS, '客户'), (K_SUP, '供应商'), (K_EXP, '费用项目'), (K_WHO, '报销人'), (K_NOTE, '备注'), (K_CHK, '校验')]
    hdr_fill = {K_IN: 'FFE2EFDA', K_OUT: 'FFFFF2CC', K_CUS: 'FFE4DFEC', K_SUP: 'FFE4DFEC', K_EXP: 'FFE4DFEC', K_WHO: 'FFE4DFEC'}
    for col, t in hdr:
        put(ws, f'{col}{CASH_HDR}', t, F_HDR_D, fill(hdr_fill.get(col, 'FFDDEBF7')), align=ACW)
    for col, t in [(K_CAT2, '实际项目'), (K_TO, '去向'), (K_CSEQ, '客户收款排序键'), (K_SSEQ, '供应商付款排序键'), (K_NET, '净额'),
                   (K_ESEQ, '费用排序键')]:
        put(ws, f'{col}{CASH_HDR}', t, F_NOTE, FILL_AUTO, align=ACW)
    ws.row_dimensions[CASH_HDR].height = 32

    # ── 数据区 ──
    rows = ctx['cash_rows']
    acc_names = ACC_NAMES
    for i in range(CASH_R1 - CASH_R0 + 1):
        r = CASH_R0 + i
        d = rows[i] if i < len(rows) else {}
        emp = f'AND({K_IN}{r}="",{K_OUT}{r}="")'
        put(ws, f'{K_SEQ}{r}', f'=IF({emp},"",ROW()-{CASH_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{K_DATE}{r}', d.get('date'), F_TXT, None, DATE, AC)
        put(ws, f'{K_ACC}{r}', d.get('account'), F_TXT, None, align=AC)
        put(ws, f'{K_CAT}{r}', d.get('category'), F_TXT, None, align=AC)
        put(ws, f'{K_MEMO}{r}', d.get('memo'), F_TXT, None, align=AL)
        put(ws, f'{K_IN}{r}', d.get('in'), F_TXT, F_IN_BG, MONEY2, AR_)
        put(ws, f'{K_OUT}{r}', d.get('out'), F_TXT, F_OUT_BG, MONEY2, AR_)
        put(ws, f'{K_ABAL}{r}',
            f'=IF(OR({K_ACC}{r}="",{emp}),"",SUMIF({acc_names},{K_ACC}{r},{SH_BASE}!${A_BAL0}${ACC_R0}:${A_BAL0}${ACC_R1})'
            f'+SUMIF({K_ACC}${CASH_R0}:{K_ACC}{r},{K_ACC}{r},{K_IN}${CASH_R0}:{K_IN}{r})'
            f'-SUMIF({K_ACC}${CASH_R0}:{K_ACC}{r},{K_ACC}{r},{K_OUT}${CASH_R0}:{K_OUT}{r}))',
            F_AUTO, FILL_AUTO, MONEY2, AR_)
        put(ws, f'{K_TBAL}{r}',
            f'=IF({emp},"",{acc0}+SUM({K_IN}${CASH_R0}:{K_IN}{r})-SUM({K_OUT}${CASH_R0}:{K_OUT}{r}))',
            F_AUTO, FILL_AUTO, MONEY2, AR_)
        put(ws, f'{K_CUS}{r}', d.get('customer'), F_TXT, None, align=AC)
        put(ws, f'{K_SUP}{r}', d.get('supplier'), F_TXT, None, align=AC)
        put(ws, f'{K_EXP}{r}', d.get('expense'), F_TXT, None, align=AC)
        put(ws, f'{K_WHO}{r}', d.get('who'), F_TXT, None, align=AC)
        put(ws, f'{K_NOTE}{r}', d.get('note'), F_TXT, None, align=AL)
        # 隐藏辅助列
        put(ws, f'{K_CAT2}{r}',
            f'=IF({K_CAT}{r}<>"",{K_CAT}{r},IF({emp},"",IF({K_CUS}{r}<>"",IF(N({K_IN}{r})>0,"客户回款",""),'
            f'IF({K_SUP}{r}<>"",IF(N({K_OUT}{r})>0,"供应商付款",""),IF({K_EXP}{r}<>"","费用支出","")))))',
            F_NOTE, border=False)
        put(ws, f'{K_TO}{r}',
            f'=IF({K_CAT2}{r}="","",IF({K_CAT}{r}="",IF({K_CUS}{r}<>"","{TO_AR}",IF({K_SUP}{r}<>"","{TO_AP}","{TO_EXP}")),'
            f'IFERROR(IF(INDEX({SH_BASE}!${C_TO}${CAT_R0}:${C_TO}${CAT_R1},'
            f'MATCH({K_CAT2}{r},{SH_BASE}!${C_NAME}${CAT_R0}:${C_NAME}${CAT_R1},0))&""="","未知",'
            f'INDEX({SH_BASE}!${C_TO}${CAT_R0}:${C_TO}${CAT_R1},'
            f'MATCH({K_CAT2}{r},{SH_BASE}!${C_NAME}${CAT_R0}:${C_NAME}${CAT_R1},0))),"未知")))', F_NOTE, border=False)
        skey = f'IF(ISNUMBER({K_DATE}{r}),{K_DATE}{r},1000000)+ROW()/100000'   # 按日期排，日期不对的排最后
        put(ws, f'{K_CSEQ}{r}', f'=IF({K_TO}{r}="{TO_AR}",{skey},"")', F_NOTE, border=False)
        put(ws, f'{K_SSEQ}{r}', f'=IF({K_TO}{r}="{TO_AP}",{skey},"")', F_NOTE, border=False)
        put(ws, f'{K_NET}{r}', f'=N({K_IN}{r})-N({K_OUT}{r})', F_NOTE, border=False)
        put(ws, f'{K_ESEQ}{r}', f'=IF({K_TO}{r}="{TO_EXP}",{skey},"")', F_NOTE, border=False)
        # 校验
        dir_ = _cat_lookup(f'{K_CAT2}{r}', C_DIR)
        chk = (f'=IF(AND({emp},{K_DATE}{r}="",{K_ACC}{r}="",{K_CAT}{r}="",{K_CUS}{r}="",{K_SUP}{r}="",{K_EXP}{r}=""),"",'
               f'IF({emp},"⚠ 没填金额（填上才进报表）",'
               f'IF(OR(AND({K_IN}{r}<>"",NOT(ISNUMBER({K_IN}{r}))),AND({K_OUT}{r}<>"",NOT(ISNUMBER({K_OUT}{r})))),"✗ 金额不是数字（可能是粘贴来的文本）",'
               f'IF(AND(N({K_IN}{r})<>0,N({K_OUT}{r})<>0),"✗ 收入、支出只能填一边",'
               f'IF(NOT(ISNUMBER({K_DATE}{r})),"✗ 日期不是真日期（要像 2026/9/25 这样填）",'
               f'IF(OR({K_DATE}{r}<年初日,{K_DATE}{r}>=年末日+1),"✗ 日期不在本会计年度",'
               f'IF({K_ACC}{r}="","✗ 没选账户",'
               f'IF(COUNTIF({acc_names},{K_ACC}{r})=0,"✗ 账户不在基础资料里",'
               f'IF(AND({K_CAT}{r}="",{K_CUS}{r}<>"",N({K_OUT}{r})>0),"✗ 付给客户的钱：收支项目选「客户退款」还是「费用支出」？",'
               f'IF(AND({K_CAT}{r}="",{K_SUP}{r}<>"",N({K_IN}{r})>0),"✗ 厂家给的钱：收支项目选「厂家返利」还是「供应商退款」？",'
               f'IF({K_CAT2}{r}="","✗ 没选收支项目",'
               f'IF({K_TO}{r}="未知","✗ 收支项目不在基础资料里",'
               f'IF(AND({K_TO}{r}="{TO_AR}",{K_CUS}{r}=""),"✗ 这笔要选客户",'
               f'IF(AND({K_TO}{r}="{TO_AR}",COUNTIF({CUS_NAMES},{K_CUS}{r})=0),"✗ 客户不在总览汇总名单里",'
               f'IF(AND({K_TO}{r}="{TO_AP}",{K_SUP}{r}=""),"✗ 这笔要选供应商",'
               f'IF(AND({K_TO}{r}="{TO_AP}",COUNTIF({SUP_NAMES},{K_SUP}{r})=0),"✗ 供应商不在基础资料里",'
               f'IF(AND({K_TO}{r}="{TO_EXP}",{K_EXP}{r}=""),"✗ 费用支出要选费用项目",'
               f'IF(AND({K_EXP}{r}<>"",COUNTIF({EXP_NAMES},{K_EXP}{r})=0),"✗ 费用项目不在基础资料里",'
               f'IF(AND({dir_}="收",N({K_OUT}{r})<>0),"✗ 这类是收款，金额填到收入列",'
               f'IF(AND({dir_}="支",N({K_IN}{r})<>0),"✗ 这类是付款，金额填到支出列",'
               f'IF(N({K_ABAL}{r})<0,"⚠ 账户余额成负数了，核对一下","√")))))))))))))))))))))')
        put(ws, f'{K_CHK}{r}', chk, F_AUTO, FILL_AUTO, align=AL)

    # 条件格式：校验列 ✗ 红、⚠ 橙
    rng_chk = f'{K_CHK}{CASH_R0}:{K_CHK}{CASH_R1}'
    ws.conditional_formatting.add(rng_chk, FormulaRule(formula=[f'LEFT(${K_CHK}{CASH_R0},1)="✗"'], fill=FILL_WARN,
                                                       font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(rng_chk, FormulaRule(formula=[f'LEFT(${K_CHK}{CASH_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng_chk, FormulaRule(formula=[f'${K_CHK}{CASH_R0}="√"'],
                                                       font=Font(name=YH, sz=10, bold=True, color='FF00B050')))

    # 下拉
    add_date_dv(ws, f'{K_DATE}{CASH_R0}:{K_DATE}{CASH_R1}')
    add_list_dv(ws, f'{K_ACC}{CASH_R0}:{K_ACC}{CASH_R1}', '=账户列表', '收/付款走的是哪个账户')
    add_list_dv(ws, f'{K_CAT}{CASH_R0}:{K_CAT}{CASH_R1}', '=收支项目列表', '选了客户/供应商/费用项目的可以不选，会自动认')
    add_list_dv(ws, f'{K_CUS}{CASH_R0}:{K_CUS}{CASH_R1}', '=客户列表', '收客户的钱选客户——自动冲减他的未付货款')
    add_list_dv(ws, f'{K_SUP}{CASH_R0}:{K_SUP}{CASH_R1}', '=供应商列表', '付厂家货款选供应商——自动冲减应付')
    add_list_dv(ws, f'{K_EXP}{CASH_R0}:{K_EXP}{CASH_R1}', '=费用项目列表', '花钱的选费用项目（就是原来费用明细表的「支出类别」）——自动进费用汇总、利润表')

    widths(ws, {K_SEQ: 6, K_DATE: 11.5, K_ACC: 11, K_CAT: 11, K_MEMO: 22, K_IN: 13, K_OUT: 13,
                K_ABAL: 15, K_TBAL: 15, K_CUS: 17, K_SUP: 14, K_EXP: 15, K_WHO: 10, K_NOTE: 17, K_CHK: 26,
                K_CAT2: 10, K_TO: 8, K_CSEQ: 6, K_SSEQ: 6, K_NET: 10, K_ESEQ: 6})
    for col in (K_CAT2, K_TO, K_CSEQ, K_SSEQ, K_NET, K_ESEQ):
        ws.column_dimensions[col].hidden = True
    ws.auto_filter.ref = f'A{CASH_HDR}:{K_CHK}{CASH_R1}'
    ws.freeze_panes = f'C{CASH_R0}'
    return ws


def build_rp(wb, ctx):
    """原【收付款明细】：保留 A:F 的样子（序号/日期/客户名称/收款金额/付款方式/备注），
    内容改成从资金台帐自动列出；右边加一块供应商付款。"""
    ws = wb[SH_RP]
    # 清掉原来手工录的数据区（数据已经整体搬进资金台帐）
    for r in range(4, max(ws.max_row, RP_R1) + 1):
        for c in range(1, 21):
            cell = ws.cell(r, c)
            cell.value = None
    if ws.auto_filter.ref:
        ws.auto_filter.ref = None
    for mr in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(mr))
    for r in (1, 2, 3):
        for c in range(1, 21):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy(ws.cell(RP_R1 + 5, 20)._style)
    widths(ws, {'A': 7, 'B': 12, 'C': 20, 'D': 15, 'E': 12, 'F': 26, 'G': 5, 'H': 3, 'I': 7, 'J': 12, 'K': 16, 'L': 15,
                'M': 12, 'N': 25, 'O': 5})
    ws.sheet_properties.tabColor = C_RP[2:]
    ws.merge_cells('A1:F1')
    put(ws, 'A1', '收付款明细台账（自动从资金台帐提取）', F_TITLE, fill(C_RP), align=AC, border=False)
    ws.merge_cells('A2:F2')
    put(ws, 'A2', '💡 不用手工登记：资金台帐里选了客户的收款自动列在这里（按日期排），冲减客户的未付货款。', F_TIP, FILL_TIP,
        align=ALW, border=False)
    header(ws, RP_HDR, [('A', '序号'), ('B', '日期'), ('C', '客户名称'), ('D', '收款金额(元)'),
                        ('E', '付款方式'), ('F', '备注')], C_RP, height=34)
    ws.row_dimensions[1].height = 33
    ws.row_dimensions[2].height = 30
    # 右边供应商块
    ws.merge_cells('I1:N1')
    put(ws, 'I1', '供应商付款明细（自动从资金台帐提取）', F_TITLE, fill(C_BUY), align=AC, border=False)
    ws.merge_cells('I2:N2')
    put(ws, 'I2', '💡 资金台帐里选了「供应商」的付款（和退款）自动列在这里，冲减欠厂家的货款。', F_TIP, FILL_TIP, align=ALW, border=False)
    header(ws, RP_HDR, [('I', '序号'), ('J', '日期'), ('K', '供应商'), ('L', '付款金额(元)'),
                        ('M', '付款方式'), ('N', '摘要')], C_BUY, height=34)
    put(ws, f'G{RP_HDR}', '行号', F_NOTE, FILL_AUTO, align=AC)
    put(ws, 'G1', f'=COUNT({SH_CASH}!${K_CSEQ}${CASH_R0}:${K_CSEQ}${CASH_R1})', F_NOTE, border=False)
    put(ws, 'O1', f'=COUNT({SH_CASH}!${K_SSEQ}${CASH_R0}:${K_SSEQ}${CASH_R1})', F_NOTE, border=False)
    put(ws, f'O{RP_HDR}', '行号', F_NOTE, FILL_AUTO, align=AC)

    idx = lambda col: f'{SH_CASH}!${col}${CASH_R0}:${col}${CASH_R1}'
    for r in range(RP_R0, RP_R1 + 1):
        k = r - RP_R0 + 1
        g = f'$G{r}'
        put(ws, f'G{r}', f'=IF({k}>$G$1,"",IFERROR(MATCH(SMALL({idx(K_CSEQ)},{k}),{idx(K_CSEQ)},0),""))', F_NOTE, border=False)
        put(ws, f'A{r}', f'=IF({g}="","",{k})', F_TXT, align=AC)
        put(ws, f'B{r}', f'=IF({g}="","",INDEX({idx(K_DATE)},{g}))', F_TXT, fmt=DATE, align=AC)
        put(ws, f'C{r}', f'=IF({g}="","",INDEX({idx(K_CUS)},{g}))', F_TXT, align=AC)
        put(ws, f'D{r}', f'=IF({g}="","",INDEX({idx(K_NET)},{g}))', F_TXT, fmt=MONEY, align=AC)
        put(ws, f'E{r}', f'=IF({g}="","",INDEX({idx(K_ACC)},{g}))', F_TXT, align=AC)
        put(ws, f'F{r}', f'=IF({g}="","",INDEX({idx(K_MEMO)},{g})&IF(INDEX({idx(K_NOTE)},{g})="","","；"&INDEX({idx(K_NOTE)},{g})))',
            F_TXT, align=AL)
        if r <= RP_SR1:
            o = f'$O{r}'
            put(ws, f'O{r}', f'=IF({k}>$O$1,"",IFERROR(MATCH(SMALL({idx(K_SSEQ)},{k}),{idx(K_SSEQ)},0),""))', F_NOTE, border=False)
            put(ws, f'I{r}', f'=IF({o}="","",{k})', F_TXT, align=AC)
            put(ws, f'J{r}', f'=IF({o}="","",INDEX({idx(K_DATE)},{o}))', F_TXT, fmt=DATE, align=AC)
            put(ws, f'K{r}', f'=IF({o}="","",INDEX({idx(K_SUP)},{o}))', F_TXT, align=AC)
            put(ws, f'L{r}', f'=IF({o}="","",-INDEX({idx(K_NET)},{o}))', F_TXT, fmt=MONEY, align=AC)
            put(ws, f'M{r}', f'=IF({o}="","",INDEX({idx(K_ACC)},{o}))', F_TXT, align=AC)
            put(ws, f'N{r}', f'=IF({o}="","",INDEX({idx(K_MEMO)},{o}))', F_TXT, align=AL)
    # 统一成 A064 的样子：白底细框（原表的淡绿/淡橙隔行底色不要了）
    style_rows(ws, RP_R0, RP_R1, 'ABCDEF', fmts={'B': DATE, 'D': MONEY}, aligns={'D': AR_, 'F': AL}, height=16)
    style_rows(ws, RP_R0, RP_SR1, 'IJKLMN', fmts={'J': DATE, 'L': MONEY}, aligns={'L': AR_, 'N': AL})
    for r in range(RP_R0, RP_R1 + 1):
        ws[f'G{r}'].font = F_NOTE
        if r <= RP_SR1:
            ws[f'O{r}'].font = F_NOTE
    ws.column_dimensions['G'].hidden = True
    ws.column_dimensions['O'].hidden = True
    ws.auto_filter.ref = f'A{RP_HDR}:F{RP_R1}'
    ws.freeze_panes = 'A4'
    return ws
