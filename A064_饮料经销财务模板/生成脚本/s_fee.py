# -*- coding: utf-8 -*-
"""补充文件（费用报销明细）那两张，现在都不用手工录，全部从【资金台帐】生成：
- 【全量费用明细总表】列跟你原来的表一模一样（序号/日期/支出类别/支出明细/报销人/支出金额/账户/备注），
  内容 = 资金台帐里所有「费用支出」按日期列出来。费用只在资金台帐记一次，费用汇总、利润表也只从资金台帐取。
- 【收款码到账对账】选月份和收款账户：每天收进收款码的钱 vs 当天从收款码转到银行的钱，差额就是手续费，
  再对一下资金台帐里记了多少手续费、收款码上还挂着多少没到账。
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *

FEE_DATE_FMT = 'yyyy.m.d'          # 跟原表一样显示 2026.6.4


def build_fee(wb, ctx):
    ws = wb.create_sheet(SH_FEE)
    ws.sheet_properties.tabColor = C_FEE[2:]
    widths(ws, {F_SEQ: 6, F_DATE: 11, F_CAT: 11, F_DETAIL: 24, F_WHO: 12, F_AMT: 12, F_ACC: 10, F_REMK: 24, F_POS: 6})
    title(ws, '全量费用明细总表（自动从资金台帐列出）', F_REMK, C_FEE,
          '💡 不用在这里录：资金台帐里记的每一笔「费用支出」（选了费用项目的）自动按日期列在这里，列跟原来的费用明细一样。'
          '「支出类别」就是资金台帐的【费用项目】，「报销人」是资金台帐的【报销人】。点表头的筛选按钮可以按月份、类别、报销人筛，'
          '上面「筛选后合计」跟着筛选变。要改哪一笔，回资金台帐改。')

    # 第 3 行：笔数 / 合计 / 筛选后合计
    ws.row_dimensions[3].height = 26
    esq = cash(K_ESEQ)
    put(ws, 'A3', '笔数：', F_KPI_L, align=AR_, border=False)
    put(ws, 'B3', f'=COUNT({esq})', F_KPI_V, fmt='0', align=AL, border=False)
    put(ws, 'C3', '金额合计：', F_KPI_L, align=AR_, border=False)
    put(ws, 'D3', f'=-SUMIFS({cash(K_NET)},{cash(K_TO)},"{TO_EXP}")', F_KPI_V, fmt=MONEY2, align=AL, border=False)
    put(ws, 'E3', '筛选后合计：', F_KPI_L, align=AR_, border=False)
    put(ws, 'F3', f'=SUBTOTAL(109,{F_AMT}{FEE_R0}:{F_AMT}{FEE_R1})', F_KPI_V, fmt=MONEY2, align=AL, border=False)
    put(ws, 'G3', f'=IF(B3>{FEE_R1 - FEE_R0 + 1},"✗ 超过 {FEE_R1 - FEE_R0 + 1} 笔列不下，跟我说一声加容量",'
                  f'"没填金额的 "&COUNTIFS({cash(K_TO)},"{TO_EXP}",{cash(K_IN)},"",{cash(K_OUT)},"")&" 笔")',
        F_NOTE, align=AL, border=False)
    ws.merge_cells('G3:H3')

    header(ws, FEE_HDR, [(F_SEQ, '序号'), (F_DATE, '日期'), (F_CAT, '支出类别'), (F_DETAIL, '支出明细'), (F_WHO, '报销人'),
                         (F_AMT, '支出金额'), (F_ACC, '账户'), (F_REMK, '备注')], C_FEE, height=30)
    put(ws, f'{F_POS}{FEE_HDR}', '行号', F_NOTE, FILL_AUTO, align=AC)

    ix = lambda col, p: f'INDEX({cash(col)},{p})'
    for r in range(FEE_R0, FEE_R1 + 1):
        k = r - FEE_R0 + 1
        p = f'${F_POS}{r}'
        ws[f'{F_POS}{r}'] = f'=IF({k}>$B$3,"",IFERROR(MATCH(SMALL({esq},{k}),{esq},0),""))'
        ws[f'{F_SEQ}{r}'] = f'=IF({p}="","",{k})'
        ws[f'{F_DATE}{r}'] = f'=IF({p}="","",IF({ix(K_DATE, p)}="","",{ix(K_DATE, p)}))'
        for col, src in ((F_CAT, K_EXP), (F_DETAIL, K_MEMO), (F_WHO, K_WHO), (F_ACC, K_ACC), (F_REMK, K_NOTE)):
            ws[f'{col}{r}'] = f'=IF({p}="","",{ix(src, p)}&"")'
        ws[f'{F_AMT}{r}'] = f'=IF({p}="","",IF(AND({ix(K_IN, p)}="",{ix(K_OUT, p)}=""),"",-{ix(K_NET, p)}))'
    style_rows(ws, FEE_R0, FEE_R1, 'ABCDEFGH', fmts={F_DATE: FEE_DATE_FMT, F_AMT: MONEY2},
               aligns={F_DETAIL: AL, F_REMK: AL, F_AMT: AR_})
    for r in range(FEE_R0, FEE_R1 + 1):
        c = ws[f'{F_POS}{r}']
        c.font = F_NOTE
    ws.column_dimensions[F_POS].hidden = True
    ws.auto_filter.ref = f'{F_SEQ}{FEE_HDR}:{F_REMK}{FEE_R1}'
    ws.freeze_panes = f'A{FEE_R0}'
    return ws


# ───────────────────────── 收款码到账对账 ─────────────────────────
QR_R0 = 6                     # 1 号在第 6 行，31 号在第 36 行
QR_TOT = QR_R0 + 31           # 37 本月合计
QR_Y_HDR = QR_TOT + 3         # 40 全年逐月表头
QR_Y0 = QR_Y_HDR + 1          # 41..52


def build_qr(wb, ctx):
    ws = wb.create_sheet(SH_QR)
    ws.sheet_properties.tabColor = C_FEE[2:]
    widths(ws, {'A': 12, 'B': 14, 'C': 14, 'D': 12, 'E': 8, 'F': 12, 'G': 12, 'H': 15, 'I': 30})
    title(ws, '收款码到账对账（自动从资金台帐生成）', 'I', C_FEE,
          '💡 不用在这里录。资金台帐里这样记，这张表自动对：① 客户扫码付款——账户选「收款码」照常记收入；'
          '② 钱到了银行——记两行「内部转账」：收款码一行支出（到账金额）、农业银行一行收入（同样金额），日期写收款那天；'
          '③ 手续费——收款码一行支出，收支项目「费用支出」，费用项目「手续费」（可以每天记，也可以月底记一笔合计）。'
          '「差额」＝收款－到账＝手续费；「还没记的手续费」不为 0 就去资金台帐补一笔。', h2=64)
    ws.row_dimensions[3].height = 28
    put(ws, 'A3', '月份：', F_KPI_L, align=AR_, border=False)
    put(ws, 'B3', ctx['kpi_month'], Font(name=YH, sz=13, bold=True, color='FF1F4E79'), FILL_SEL, '0"月"', AC)
    add_list_dv(ws, 'B3', '"1,2,3,4,5,6,7,8,9,10,11,12"', '选月份')
    put(ws, 'C3', '收款账户：', F_KPI_L, align=AR_, border=False)
    put(ws, 'D3', '收款码', Font(name=YH, sz=12, bold=True, color='FF1F4E79'), FILL_SEL, align=AC)
    add_list_dv(ws, 'D3', '=账户列表', '对哪个收款账户（默认收款码；别的收款渠道也能对）')
    put(ws, 'E3', '=IF(COUNTIF(账户列表,D3)=0,"✗ 账户不在基础资料里","")', F_RED, align=AL, border=False)
    ws.merge_cells('E3:G3')

    hdr = [('A', '日期'), ('B', '收款\n(当天收进来的)'), ('C', '到账\n(当天转去银行的)'), ('D', '差额\n(收款－到账)'),
           ('E', '费率'), ('F', '已记手续费'), ('G', '还没记的\n手续费'), ('H', '账户余额\n(当天结束)'), ('I', '状态')]
    header(ws, 5, hdr, C_FEE, height=36)

    acc, M = '$D$3', '$B$3'
    bal0 = f'SUMIF({ACC_NAMES},{acc},{SH_BASE}!${A_BAL0}${ACC_R0}:${A_BAL0}${ACC_R1})'

    def day_sum(col, to_crit, d0, d1):
        return (f'SUMIFS({cash(col)},{cash(K_ACC)},{acc},{cash(K_TO)},"{to_crit}",'
                f'{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&{d1})')

    for i in range(31):
        r = QR_R0 + i
        d = f'$A{r}'
        ws[f'A{r}'] = f'=IF(DAY(DATE(年度,{M},{i + 1}))={i + 1},DATE(年度,{M},{i + 1}),"")'
        rows_f = {
            'B': f'=IF({d}="","",{day_sum(K_IN, "<>" + TO_XFER, d, d + "+1")})',
            'C': f'=IF({d}="","",{day_sum(K_OUT, TO_XFER, d, d + "+1")})',
            'D': f'=IF({d}="","",B{r}-C{r})',
            'E': f'=IF(OR({d}="",N(B{r})=0,N(C{r})=0),"",D{r}/B{r})',
            'F': f'=IF({d}="","",{day_sum(K_OUT, TO_EXP, d, d + "+1")})',
            'G': f'=IF(OR({d}="",N(C{r})=0),"",D{r}-F{r})',
            'H': f'=IF({d}="","",{bal0}+SUMIFS({cash(K_NET)},{cash(K_ACC)},{acc},{cash(K_DATE)},"<"&({d}+1)))',
            'I': (f'=IF({d}="","",IF(AND(N(B{r})=0,N(C{r})=0,N(F{r})=0),"",IF(N(C{r})=0,"⏳ 还没记到账",'
                  f'IF(ROUND(D{r}-F{r},2)=0,"√ 对上了",IF(D{r}-F{r}>0,"⚠ 差 "&TEXT(D{r}-F{r},"0.00")&" 元，手续费没记？",'
                  f'"⚠ 到账＋手续费比收款多 "&TEXT(F{r}-D{r},"0.00")&" 元")))))'),
        }
        for col, f in rows_f.items():
            ws[f'{col}{r}'] = f
    style_rows(ws, QR_R0, QR_R0 + 30, 'ABCDEFGHI',
               fmts={'A': 'm"月"d"日"', 'B': MONEY2, 'C': MONEY2, 'D': MONEY2, 'E': '0.00%', 'F': MONEY2,
                     'G': MONEY2, 'H': MONEY2},
               aligns={c: AR_ for c in 'BCDEFGH'} | {'I': AL}, bold=('D',))
    # 本月合计
    r = QR_TOT
    put(ws, f'A{r}', '本月合计', F_TXTB, FILL_TOT, align=AC)
    for col in 'BCDFG':
        put(ws, f'{col}{r}', f'=SUM({col}{QR_R0}:{col}{QR_TOT - 1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    # 费率只算已经到账的那几天（没到账的收款不算进去，不然费率虚高）
    put(ws, f'E{r}', (f'=IF(SUMIF(C{QR_R0}:C{QR_TOT - 1},">0",B{QR_R0}:B{QR_TOT - 1})=0,"",'
                      f'SUMIF(C{QR_R0}:C{QR_TOT - 1},">0",D{QR_R0}:D{QR_TOT - 1})/SUMIF(C{QR_R0}:C{QR_TOT - 1},">0",B{QR_R0}:B{QR_TOT - 1}))'),
        F_TXTB, FILL_TOT, '0.00%', AR_)
    put(ws, f'H{r}', f'={bal0}+SUMIFS({cash(K_NET)},{cash(K_ACC)},{acc},{cash(K_DATE)},"<"&DATE(年度,{M}+1,1))',
        F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'I{r}', f'=IF(ROUND(G{r},2)=0,"","⚠ 本月还有 "&TEXT(G{r},"#,##0.00")&" 元手续费没记")', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'A{r + 1}', '「账户余额」＝到这天为止收款码上还挂着的钱（还没到账的＋还没记的手续费），月底一般只剩最后一两天还没到账的收款。',
        F_NOTE, align=AL, border=False)
    ws.merge_cells(f'A{r + 1}:I{r + 1}')

    # 全年逐月
    put(ws, f'A{QR_Y_HDR - 1}', '全年逐月（同一个收款账户）', F_SEC, border=False)
    header(ws, QR_Y_HDR, [('A', '月份'), ('B', '收款'), ('C', '到账'), ('D', '已记手续费'), ('E', '费率'),
                          ('F', '本月新挂账\n(收款－到账－手续费)'), ('G', '月末账户余额'), ('H', ''), ('I', '')], C_FEE, height=34)
    for m in range(1, 13):
        r = QR_Y0 + m - 1
        d0, d1 = f'DATE(年度,{m},1)', f'DATE(年度,{m}+1,1)'
        put(ws, f'A{r}', m, fmt='0"月"')
        ws[f'B{r}'] = f'={day_sum(K_IN, "<>" + TO_XFER, d0, d1)}'
        ws[f'C{r}'] = f'={day_sum(K_OUT, TO_XFER, d0, d1)}'
        ws[f'D{r}'] = f'={day_sum(K_OUT, TO_EXP, d0, d1)}'
        ws[f'E{r}'] = f'=IF(N(C{r})+N(D{r})=0,"",D{r}/(C{r}+D{r}))'       # 手续费 ÷ 已结算的收款（到账＋手续费）
        ws[f'F{r}'] = f'=B{r}-C{r}-D{r}'
        ws[f'G{r}'] = f'={bal0}+SUMIFS({cash(K_NET)},{cash(K_ACC)},{acc},{cash(K_DATE)},"<"&{d1})'
    style_rows(ws, QR_Y0, QR_Y0 + 11, 'ABCDEFG',
               fmts={'A': '0"月"', 'B': MONEY2, 'C': MONEY2, 'D': MONEY2, 'E': '0.00%', 'F': MONEY2, 'G': MONEY2},
               aligns={c: AR_ for c in 'BCDEFG'})
    r = QR_Y0 + 12
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for col in 'BCDF':
        put(ws, f'{col}{r}', f'=SUM({col}{QR_Y0}:{col}{QR_Y0 + 11})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'E{r}', f'=IF(N(C{r})+N(D{r})=0,"",D{r}/(C{r}+D{r}))', F_TXTB, FILL_TOT, '0.00%', AR_)
    put(ws, f'G{r}', f'=G{QR_Y0 + 11}', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'A{r + 2}', '你原来那张 6 月对账表（收款码 vs 农业银行，6 月手续费合计 952.09）原样放在 参考/原表3 里；'
                         '它的逐日收款是按收款码流水记的，资金台帐里是按客户记的收款，两边对不上号，所以没有搬进资金台帐。',
        F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'A{r + 2}:I{r + 3}')

    rng = f'I{QR_R0}:I{QR_R0 + 30}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($I{QR_R0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT($I{QR_R0},1)="√"'],
                                                   font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    ws.freeze_panes = f'B{QR_R0}'
    return ws
