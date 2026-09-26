# -*- coding: utf-8 -*-
"""独立复算校验：不看表里的公式，直接从两个原表（套上 fixes.py 写明的修正）用 Python 重算关键数，
再跟 LibreOffice 在副本上算出来的值（_lo_copy.xlsx，make_all.py 会留一份）逐项比。

跑法：python3 verify_a064.py [_lo_copy.xlsx]
返回码：0 全部通过，1 有对不上的
"""
import os
import sys
import datetime as dt
from collections import defaultdict, OrderedDict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import fixes
import data_prep
from common import *
from build_a064 import SRC1, SRC2

TOL = 0.01
fails = []


def ok(name, got, exp, tol=TOL):
    good = (got is None and exp is None) or (isinstance(got, (int, float)) and isinstance(exp, (int, float))
                                              and abs(got - exp) <= tol) or got == exp
    print(f'  {"✓" if good else "✗"} {name}：表里 {got!r}，复算 {exp!r}')
    if not good:
        fails.append(name)


def num(v):
    return v if isinstance(v, (int, float)) and not isinstance(v, bool) else 0.0


def main(lo_path):
    s1 = data_prep.Src1(SRC1)
    wv = openpyxl.load_workbook(lo_path, data_only=True)
    Y = fixes.BOOK_YEAR

    # ── 原始数据（套上修正）──
    outs = s1.outbound()
    deps = s1.deposits()
    recs = s1.receipts()
    out_date = {r['row']: (fixes.OUT_DATE_FIX.get(r['row']).date() if r['row'] in fixes.OUT_DATE_FIX else r['date'])
                for r in outs}
    amt = lambda r: num(r['amount_cached'])       # 金额列：公式按缓存值，手填按手填
    print('【1】总量对账')
    tot_out = sum(amt(r) for r in outs)
    tot_rec = sum(num(r['amount']) for r in recs)
    tot_dep = sum(amt(r) for r in deps)
    ov = wv[SH_OV]
    ok('出库明细金额合计 = 总览汇总已领用合计', ov[f'D{OV_TOT}'].value, tot_out)
    ok('原收付款明细合计 = 总览汇总已付款合计', ov[f'F{OV_TOT}'].value, tot_rec)
    cash_ws = wv[SH_CASH]
    ok('原收付款明细合计 = 资金台帐收入合计',
       sum(num(cash_ws[f'{K_IN}{r}'].value) for r in range(CASH_R0, CASH_R1 + 1)), tot_rec)
    ok('资金台帐笔数 = 原收付款明细笔数',
       sum(1 for r in range(CASH_R0, CASH_R1 + 1) if cash_ws[f'{K_IN}{r}'].value not in (None, '')), len(recs))
    rp = wv[SH_RP]
    ok('收付款明细（自动提取）合计', sum(num(rp[f'D{r}'].value) for r in range(RP_R0, RP_R1 + 1)), tot_rec)
    ok('存条明细金额合计 = 商品库存累计入库金额合计', wv[SH_INV][f'F{INV_TOT}'].value, tot_dep)
    ok('出库明细金额合计 = 商品库存已领用金额合计', wv[SH_INV][f'H{INV_TOT}'].value, tot_out)

    print('【2】按客户：已领用 / 已付款 / 未付货款')
    cust = data_prep.build_customer_master(s1)
    by_out, by_rec = defaultdict(float), defaultdict(float)
    for r in outs:
        by_out[r['customer']] += amt(r)
    for r in recs:
        by_rec[r['customer']] += num(r['amount'])
    bad = 0
    for i, (name, dep, _) in enumerate(cust):
        row = OV_R0 + i
        if ov[f'B{row}'].value != name:
            bad += 1
            print(f'    ✗ 第 {row} 行客户应为 {name}，表里是 {ov[f"B{row}"].value}')
            continue
        for col, exp in (('D', by_out[name]), ('F', by_rec[name]), ('G', by_out[name] - by_rec[name])):
            got = num(ov[f'{col}{row}'].value)
            if abs(got - exp) > TOL:
                bad += 1
                print(f'    ✗ {name} {col}列：表里 {got}，复算 {exp}')
    ok(f'{len(cust)} 个客户逐个核对，不一致的格数', bad, 0)

    print('【3】商品库存：「客户+商品」组合')
    combos = OrderedDict()
    for r in deps + outs:
        if r['customer'] and r['goods']:
            combos.setdefault((r['customer'], r['goods']), None)
    inv = wv[SH_INV]
    got = {(inv[f'B{r}'].value, inv[f'C{r}'].value) for r in range(INV_R0, INV_R1 + 1) if inv[f'B{r}'].value}
    ok('组合个数', len(got), len(combos))
    ok('组合完全一致（漏的+多的）', len(set(combos) ^ got), 0)

    print('【4】按月销售收入（利润表）')
    pl = wv[SH_PL]
    plrow = {str(pl[f'A{r}'].value).strip(): r for r in range(6, 40) if pl[f'A{r}'].value}
    R_REV, R_LOSS, R_NP = plrow['主营业务收入（出库领用）'], plrow['商品报损损失'], plrow['五、净利润']
    by_m = defaultdict(float)
    for r in outs:
        d = out_date[r['row']]
        if d and d.year == Y:
            by_m[d.month] += amt(r)
    for m in range(1, 13):
        col = CL(4 + m)
        if by_m[m] or num(pl[f'{col}{R_REV}'].value):
            ok(f'{m} 月主营业务收入', pl[f'{col}{R_REV}'].value, by_m[m])

    print('【5】报损')
    s2 = openpyxl.load_workbook(SRC2, data_only=False)
    month_sheet = fixes.BAOSUN_MONTHS
    total = 0.0
    for m, sh in month_sheet.items():
        if sh == SH_BS9:
            continue
        ws = s2[sh]
        s = 0.0
        for r in range(6, ws.max_row + 1):
            if ws.cell(r, 2).value and ws.cell(r, 3).value:
                e = ws.cell(r, 5).value
                if isinstance(e, str) and e.startswith('='):
                    e = eval(e[1:])
                g = ws.cell(r, 7).value
                if isinstance(e, (int, float)) and isinstance(g, (int, float)):
                    s += e * g
                elif e not in (None, '') and g in (None, ''):
                    pass                     # 有数量没单价：金额按 0（【数据校验】会提示）
        total += s
        ok(f'{m} 月报损（{sh}）', num(pl[f'{CL(4 + m)}{R_LOSS}'].value), s)
    ok('报损汇总一览 全年总金额', wv[SH_BSS]['B3'].value, total)

    print('【6】资产负债表平衡 & 报错')
    bal = wv[SH_BAL]
    rA = next(r for r in range(6, 40) if bal[f'A{r}'].value == '资产总计')
    rN = next(r for r in range(6, 40) if str(bal[f'F{r}'].value).startswith('加：本年利润'))
    ok('资产总计 − 负债和所有者权益合计', round(num(bal[f'C{rA}'].value) - num(bal[f'H{rA}'].value), 2), 0)
    ok('本年利润（资产负债表）= 利润表本年累计净利润（同一月份时）',
       bal[f'H{rN}'].value if bal['B3'].value == pl['B3'].value else None,
       pl[f'D{R_NP}'].value if bal['B3'].value == pl['B3'].value else None)
    errs = 0
    for ws in wv.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith('#') and c.value.rstrip('!?0/').upper() in (
                        '#N/A', '#VALUE', '#REF', '#DIV', '#NUM', '#NAME', '#NULL'):
                    errs += 1
    ok('全册公式报错格数', errs, 0)

    print()
    if fails:
        print(f'✗ {len(fails)} 项对不上：{fails}')
        return 1
    print('✓ 全部核对通过')
    return 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '_lo_copy.xlsx')))
