# -*- coding: utf-8 -*-
"""【数据校验】一张表把所有「会让报表算错」的情况数出来；【首页】导航 + 关键数。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *
from s_reports import PL, BAL, STK_R0, STK_R1, STK_TOT, PL_MCOL
from s_baosun import AUX_BS_UNM, bsx_amt

AUX_CHK_OUT = 'N'     # _辅助 N 列：出库明细每行 金额≠数量×单价 标 1
AUX_CHK_CUN = 'O'     # _辅助 O 列：存条明细


def build_aux_check(wb, ctx):
    ws = wb[SH_AUX]
    put(ws, f'{AUX_CHK_OUT}3', '出库:金额≠数量×单价', F_NOTE, border=False)
    put(ws, f'{AUX_CHK_CUN}3', '存条:金额≠数量×单价', F_NOTE, border=False)
    for s in range(OUT_R0, OUT_R1 + 1):
        ws[f'{AUX_CHK_OUT}{s}'] = (f'=IF(AND(ISNUMBER({SH_OUT}!G{s}),ISNUMBER({SH_OUT}!E{s}),ISNUMBER({SH_OUT}!F{s})),'
                                   f'IF(ABS({SH_OUT}!G{s}-{SH_OUT}!E{s}*{SH_OUT}!F{s})>0.005,1,0),0)')
    for s in range(CUN_R0, CUN_R1 + 1):
        ws[f'{AUX_CHK_CUN}{s}'] = (f'=IF(AND(ISNUMBER({SH_CUN}!G{s}),ISNUMBER({SH_CUN}!E{s}),ISNUMBER({SH_CUN}!F{s})),'
                                   f'IF(ABS({SH_CUN}!G{s}-{SH_CUN}!E{s}*{SH_CUN}!F{s})>0.005,1,0),0)')


def build_check(wb, ctx):
    ws = wb.create_sheet(SH_CHK)
    ws.sheet_properties.tabColor = C_CHK[2:]
    title(ws, '数据校验（全自动 · 每项应为 √）', 'F', C_CHK,
          '💡 这里把「会让报表算错」的情况一项项数出来。✗ 要改；⚠ 是提醒，看一眼确认没问题就行。改完回来看，变成 √ 就好了。')
    heads = ['序号', '检查项', '数量/金额', '状态', '怎么处理', '去哪看']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}4', t, F_HDR, fill(C_CHK), align=ACW)
    ws.row_dimensions[4].height = 28
    cusn = CUS_NAMES
    oc, od, ob = out(O_CUS), out(O_GOODS), out(O_DATE)
    uc, ud, ub = cun(U_CUS), cun(U_GOODS), cun(U_DATE)
    stk = lambda col: f'{SH_STK}!${col}${STK_R0}:${col}${STK_R1}'
    checks = [
        # (检查项, 公式, 级别 'E' 错 / 'W' 提醒, 怎么处理, 去哪)
        ('资金台帐：校验列打 ✗ 的行', f'=COUNTIF({cash(K_CHK)},"✗*")', 'E', '筛选校验列，按提示改', SH_CASH),
        ('资金台帐：账户余额变成负数的行', f'=COUNTIF({cash(K_CHK)},"⚠*")', 'W', '多半是漏记了一笔收入，或账户选错', SH_CASH),
        ('出库明细：客户不在总览汇总名单里', f'=SUMPRODUCT(({oc}<>"")*(COUNTIF({cusn},{oc})=0))', 'E',
         '去【总览汇总】B 列加上这个客户，或者把名字改成名单里的写法', SH_OUT),
        ('存条明细：客户不在总览汇总名单里', f'=SUMPRODUCT(({uc}<>"")*(COUNTIF({cusn},{uc})=0))', 'E', '同上', SH_CUN),
        ('出库明细：商品不在商品档案里', f'=SUMPRODUCT(({od}<>"")*(COUNTIF({GOODS_NAMES},{od})=0))', 'E',
         '去【基础资料·商品档案】加上，或改成档案里的写法', SH_OUT),
        ('存条明细：商品不在商品档案里', f'=SUMPRODUCT(({ud}<>"")*(COUNTIF({GOODS_NAMES},{ud})=0))', 'E', '同上', SH_CUN),
        ('出库明细：日期不是真日期或不在本年度',
         f'=COUNTIF({oc},"?*")-COUNTIFS({oc},"?*",{ob},">="&年初日,{ob},"<="&年末日)', 'E', '日期要像 2026/9/25 这样录', SH_OUT),
        ('存条明细：日期不是真日期或不在本年度',
         f'=COUNTIF({uc},"?*")-COUNTIFS({uc},"?*",{ub},">="&年初日,{ub},"<="&年末日)', 'E', '同上', SH_CUN),
        ('出库明细：金额≠数量×单价（金额被手工改过）', f'=SUM({SH_AUX}!{AUX_CHK_OUT}{OUT_R0}:{AUX_CHK_OUT}{OUT_R1})', 'W',
         '金额列本来是公式，被手填了数；确认一下是不是故意的', SH_OUT),
        ('存条明细：金额≠数量×单价（金额被手工改过）', f'=SUM({SH_AUX}!{AUX_CHK_CUN}{CUN_R0}:{AUX_CHK_CUN}{CUN_R1})', 'W', '同上', SH_CUN),
        ('出库明细：有数量没单价的行（赠品）', f'=COUNTIFS({oc},"?*",{out(O_QTY)},"<>",{out(O_PRICE)},"")', 'W',
         '赠品这样录没问题：收入为 0，成本照算', SH_OUT),
        ('采购进货：校验列打 ✗ 的行', f'=COUNTIF({buy(B_CHK)},"✗*")', 'E', '按校验列提示改', SH_BUY),
        ('公司库存：结存为负的商品', f'=COUNTIF({stk("Q")},"✗*")', 'E', '进货没录全：把进货单补进【采购进货】，或在商品档案填期初库存', SH_STK),
        ('公司库存：有出库却没有进价的商品', f'=COUNTIF({stk("Q")},"*没有进价*")', 'W',
         '在【基础资料·商品档案】填「参考进价」，不然这些商品成本按 0 算、毛利虚高', SH_STK),
        ('公司库存：仓库的货不够兑现客户存条', f'=COUNTIF({stk("Q")},"*不够兑现*")', 'W', '该进货了（或进货没录）', SH_STK),
        ('总览汇总：客户名字重复', f'=SUMPRODUCT(({cusn}<>"")*(COUNTIF({cusn},{cusn})>1))', 'E',
         '同一个客户只能占一行，不然领用和收款会算两遍', SH_OV),
        ('商品档案：商品名字重复', f'=SUMPRODUCT(({GOODS_NAMES}<>"")*(COUNTIF({GOODS_NAMES},{GOODS_NAMES})>1))', 'E', '删掉重复的那一行', SH_BASE),
        ('资金台帐选的客户，收款没进总览汇总的金额',
         f'=ROUND(SUMIFS({cash(K_NET)},{cash(K_TO)},"{TO_AR}")-{SH_OV}!F{OV_TOT},2)', 'E',
         '客户名没在总览汇总登记、或者没选客户；【资产负债表】附注也有这个数', SH_CASH),
        ('出库明细的金额，没进总览汇总的部分',
         f'=ROUND(SUM({out(O_AMT)})-{SH_OV}!D{OV_TOT},2)', 'E', '有出库的客户没在总览汇总登记', SH_OV),
        ('商品库存合计 ≠ 存条明细合计（有行没填客户或商品）',
         f'=ROUND(SUM({cun(U_AMT)})-{SH_INV}!F{INV_TOT},2)', 'E', '存条明细里有行没填客户或商品', SH_CUN),
        ('资产负债表：不平的金额', f'={SH_BAL}!C18', 'E', '看资产负债表附注', SH_BAL),
        ('资金台帐：内部转账没配对（转出≠转入）', f'={SH_BAL}!C28', 'E', '内部转账要记一出一进两行', SH_CASH),
        ('资金台帐：没分类的收支净额', f'={SH_BAL}!C27', 'E', '在资金台帐里选上收支项目', SH_CASH),
        ('报损表里没对上商品档案的报损金额（全年）',
         f'=ROUND(SUM({SH_AUX}!{bsx_amt(1)}{AUX_BS_UNM}:{bsx_amt(12)}{AUX_BS_UNM}),2)', 'W',
         '在商品档案「报损表品名」填上报损表里的写法（不影响利润，只影响库存按商品拆分）', SH_BASE),
        ('报损月表清单：填了但找不到的表名', f'=COUNTIF({SH_BASE}!{M_NOTE}{BSM_R0}:{M_NOTE}{BSM_R1},"✗*")', 'E',
         '表名要和工作表标签一字不差', SH_BASE),
        ('资金账户：当前余额为负的账户', f'=COUNTIF({SH_BASE}!{A_BALNOW}{ACC_R0}:{A_BALNOW}{ACC_R1},"<0")', 'W',
         '年初余额没填，或漏记了收入', SH_BASE),
    ]
    for i, (t, f, lvl, how, where) in enumerate(checks):
        r = 5 + i
        put(ws, f'A{r}', i + 1, F_AUTO, align=AC)
        put(ws, f'B{r}', t, F_TXT, align=ALW)
        put(ws, f'C{r}', f, F_TXTB, fmt='#,##0.##;[Red]-#,##0.##;0', align=AC)
        bad = '✗ 要改' if lvl == 'E' else '⚠ 看一下'
        put(ws, f'D{r}', f'=IF(ROUND(N(C{r}),2)=0,"√","{bad}")', F_TXTB, align=AC)
        put(ws, f'E{r}', how, F_NOTE, align=ALW)
        c = put(ws, f'F{r}', f'→ {where}', Font(name=YH, sz=10, color='FF0563C1', underline='single'), align=AL)
        c.hyperlink = f"#'{where}'!A1"
        ws.row_dimensions[r].height = 30
    last = 5 + len(checks) - 1
    ws.conditional_formatting.add(f'D5:D{last}', FormulaRule(formula=['LEFT($D5,1)="✗"'], fill=FILL_WARN,
                                                            font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'D5:D{last}', FormulaRule(formula=['LEFT($D5,1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'D5:D{last}', FormulaRule(formula=['$D5="√"'], fill=FILL_OK,
                                                            font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    r = last + 2
    put(ws, f'B{r}', '要改（✗）的项数', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'C{r}', f'=COUNTIF(D5:D{last},"✗*")', F_RED, FILL_TOT, align=AC)
    put(ws, f'B{r + 1}', '提醒（⚠）的项数', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'C{r + 1}', f'=COUNTIF(D5:D{last},"⚠*")', F_TXTB, FILL_TOT, align=AC)
    ctx['chk_err_cell'] = f'{SH_CHK}!$C${r}'
    ctx['chk_warn_cell'] = f'{SH_CHK}!$C${r + 1}'
    widths(ws, {'A': 5, 'B': 38, 'C': 13, 'D': 11, 'E': 46, 'F': 14})
    ws.freeze_panes = 'A5'
    return ws


def build_home(wb, ctx):
    ws = wb.create_sheet(SH_HOME, 0)
    ws.sheet_properties.tabColor = 'C00000'
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:H1')
    put(ws, 'A1', '=IF(基础资料!B6="","饮料经销财务模板",基础资料!B6)&" · "&年度&"年度账"', Font(name=YH, sz=20, bold=True, color='FFFFFFFF'),
        fill('FFC00000'), align=AC, border=False)
    ws.row_dimensions[1].height = 44
    ws.merge_cells('A2:H2')
    put(ws, 'A2', '💡 日常只录四张表：资金台帐（收钱付钱）、出库明细（送货）、存条明细（客户存条）、采购进货（进货）；报损照原来每月一张。'
                  '其余全部自动：总览、库存、收付款明细、对账单、利润表、资产负债表。淡黄/白格子手填，灰格子别动。',
        F_TIP, FILL_TIP, align=ALW, border=False)
    ws.row_dimensions[2].height = 36

    # 关键数
    put(ws, 'A4', '关键数（截至今天录入的数据）', F_SEC, border=False)
    np_row, rev_row = PL['np'], PL['rev_main']
    kpis = [
        ('本年销售额', f'={SH_PL}!Q{rev_row}'),
        ('本年毛利', f'={SH_PL}!Q{PL["gp"]}'),
        ('本年净利润', f'={SH_PL}!Q{np_row}'),
        ('当前资金余额', f'={SH_BASE}!{A_BALNOW}{ACC_R1 + 1}'),
        ('客户欠我们', f'=SUMIF({SH_OV}!$G${OV_R0}:$G${OV_R1},">0")'),
        ('客户存条/预存', f'=-SUMIF({SH_OV}!$G${OV_R0}:$G${OV_R1},"<0")'),
        ('欠厂家货款', f'=SUMIF({SH_AUX}!$K${SUP_R0}:$K${SUP_R1},">0")'),
        ('数据校验待改项', f'={ctx["chk_err_cell"]}'),
    ]
    for i, (lab, f) in enumerate(kpis):
        col = CL(i + 1)
        put(ws, f'{col}5', lab, Font(name=YH, sz=10, bold=True, color='FF595959'), fill('FFF2F2F2'), align=AC)
        put(ws, f'{col}6', f, Font(name=YH, sz=14, bold=True, color='FFC00000' if i < 7 else 'FF1F4E79'),
            fill('FFF2F2F2'), MONEY2 if i < 7 else '0"项"', AC)
        ws.column_dimensions[col].width = 17
    ws.row_dimensions[6].height = 32
    put(ws, 'A7', '「欠厂家货款」按资产负债表选定月份月底算；其余是全年/当前数。', F_NOTE, border=False)
    ws.merge_cells('A7:H7')

    # 导航
    put(ws, 'A9', '表格导航（点表名直接跳过去）', F_SEC, border=False)
    for i, t in enumerate(['表名', '干什么的', '你要做什么']):
        put(ws, f'{"ABD"[i]}10', t, F_HDR, fill('FF595959'), align=AC)
    ws.merge_cells('B10:C10')
    ws.merge_cells('D10:H10')
    nav = [
        (SH_CASH, '收支混合录入，逐笔出账户余额和总余额', '★ 每收/付一笔钱记一行；收客户的选客户，付厂家的选供应商，花钱的选费用项目'),
        (SH_OUT, '（原表）客户领货/送货', '★ 每次送货记一行'),
        (SH_CUN, '（原表）客户存条（预付的货）', '★ 客户存条时记，一个商品一行'),
        (SH_BUY, '（新）从厂家进货', '★ 每张进货单记，厂家搭赠填赠品数量'),
        (SH_BS9, '（原表）每月过期报损', '★ 每月一张，照原来的填；新月份复制上月的表，并在基础资料登记表名'),
        (SH_BASE, '商品、供应商、资金账户、收支项目、费用项目', '新商品/新账户先在这里加；商品的参考进价要填'),
        (SH_OV, '（原表）客户存条&货款总览——客户名单也在这', '新客户在 B 列加；已付款自动从资金台帐取'),
        (SH_INV, '（原表）每个客户每个商品的存条库存', '全自动'),
        (SH_RP, '（原表）收付款明细', '全自动，从资金台帐提取'),
        (SH_Q, '（原表）客户快速查询', '选客户看'),
        (SH_RPS, '按月收付、按客户/供应商汇总', '全自动'),
        (SH_CST, '给客户的对账单（可打印签字）', '选客户、起止日期'),
        (SH_SST, '和厂家的对账单', '选供应商、起止日期'),
        (SH_STK, '公司仓库进销存（按商品）', '选月份看'),
        (SH_BSS, '（原表）过期报损汇总', '选全年或某月看'),
        (SH_EXP, '费用按项目×月份', '全自动'),
        (SH_PL, '利润表（本月 / 本年累计 / 逐月）', '选月份看'),
        (SH_BAL, '资产负债表（简易）', '选月份看，平衡检查要是 √'),
        (SH_CHK, '所有数据问题一表列出', '有 ✗ 就按提示改'),
    ]
    for i, (sh, what, todo) in enumerate(nav):
        r = 11 + i
        c = put(ws, f'A{r}', sh, Font(name=YH, sz=11, bold=True, color='FF0563C1', underline='single'), align=AC)
        c.hyperlink = f"#'{sh}'!A1"
        ws.merge_cells(f'B{r}:C{r}')
        put(ws, f'B{r}', what, F_TXT, align=ALW)
        put(ws, f'C{r}', None)
        ws.merge_cells(f'D{r}:H{r}')
        put(ws, f'D{r}', todo, F_IN if todo.startswith('★') else F_NOTE, FILL_IN if todo.startswith('★') else None, align=ALW)
        for cc in 'EFGH':
            put(ws, f'{cc}{r}', None, fill_=FILL_IN if todo.startswith('★') else None)
        ws.row_dimensions[r].height = 24
    r = 11 + len(nav) + 1
    put(ws, f'A{r}', '钱和货是怎么流的', F_SEC, border=False)
    flow = [
        '客户付存条款/货款 → 资金台帐（选客户）→ 总览汇总「已付款」、收付款明细、客户对账单',
        '客户领货 → 出库明细 → 总览汇总「已领用」、商品库存、利润表「主营业务收入」、公司库存「出库」',
        '从厂家进货 → 采购进货 → 公司库存、成本计算（加权平均进价）、供应商对账单「进货」',
        '付厂家货款 → 资金台帐（选供应商）→ 供应商对账单「付款」、资产负债表「应付账款」',
        '各项开支 → 资金台帐（选费用项目）→ 费用汇总 → 利润表',
        '过期报损 → 各月报损表 → 报损汇总一览、利润表「商品报损损失」、公司库存「报损」',
    ]
    for i, t in enumerate(flow):
        ws.merge_cells(f'A{r + 1 + i}:H{r + 1 + i}')
        put(ws, f'A{r + 1 + i}', f'{i + 1}. {t}', F_TXT, align=AL, border=False)
    r = r + len(flow) + 2
    put(ws, f'A{r}', '颜色', F_SEC, border=False)
    legend = [(FILL_IN, '淡黄：手填的格子'), (FILL_AUTO, '淡灰：公式自动算，别改'), (FILL_SEL, '亮黄：查询/报表的选择格'),
              (FILL_WARN, '淡红：校验发现的问题')]
    for i, (fl, t) in enumerate(legend):
        put(ws, f'{CL(i * 2 + 1)}{r + 1}', None, fill_=fl)
        put(ws, f'{CL(i * 2 + 2)}{r + 1}', t, F_NOTE, align=AL, border=False)
    return ws
