# -*- coding: utf-8 -*-
"""原表 2（过期报损）：四张表连格式原样搬进来，只修公式；照 8 月的样子加一张空白的 9 月表；
【报损汇总一览】保留原来的版式，改成按【基础资料·报损月表清单】汇总所有月份，
不再用 UNIQUE / XLOOKUP（老版 WPS、Excel 2016 都认不了）。

报损数据的「中转站」放在 _辅助 表右边（每个商品每个月的报损金额/数量），
成本计算、利润表、报损汇总一览都从中转站取，INDIRECT 只在中转站里出现。"""
from copy import copy
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

from common import *

BS_ROWS = (6, 1000)          # 各月报损表的数据行范围

# ── _辅助 里报损中转站的位置 ──
AUX_BS_P0 = 6                                  # 商品块：第 6..105 行 ↔ 基础资料 5..104
AUX_BS_P1 = AUX_BS_P0 + (BASE_R1 - BASE_R0)
AUX_BS_S0 = 110                                # 供应商块：第 110..159 行 ↔ 基础资料 Q5..Q54
AUX_BS_S1 = AUX_BS_S0 + (SUP_R1 - SUP_R0)
AUX_BS_TOT = 162                               # 每月报损总金额
AUX_BS_TOTQ = 163                              # 每月报损总数量（原始数，瓶/罐）
AUX_BS_UNM = 164                               # 每月没对上商品档案的报损金额
AUX_BS_NOPRICE = 165                           # 每月「有数量没单价」的行数
AUX_BS_NOCOMP = 166                            # 每月「有产品没填所属公司」的行数
AUX_BS_NONNUM = 167                            # 每月数量/单价不是数字的格数
AUX_BS_NOH = 168                               # 每月「数量、单价都有，金额格却没公式」的行数（中间插了行）
BSX_NAME, BSX_PACK = 'AA', 'AB'
BSX_AMT0 = CI('AC')          # 金额 12 列 AC..AN
BSX_RAW0 = BSX_AMT0 + 12     # 原始数量 AO..AZ
BSX_BOX0 = BSX_AMT0 + 24     # 按箱/件录的数量 BA..BL
BSX_PCS0 = BSX_AMT0 + 36     # 折成件 BM..BX
BSX_SEL = CL(BSX_AMT0 + 48)  # BY 选定月份金额
BSX_KEY = CL(BSX_AMT0 + 49)  # BZ 排名键
BSX_SELQ = CL(BSX_AMT0 + 50) # CA 选定月份数量


def bsx_amt(m):  return CL(BSX_AMT0 + m - 1)
def bsx_pcs(m):  return CL(BSX_PCS0 + m - 1)


def _copy_sheet(src_ws, dst_ws):
    """跨工作簿复制：值、样式、列宽、行高、合并、冻结、视图、打印设置"""
    for row in src_ws.iter_rows():
        for c in row:
            d = dst_ws.cell(c.row, c.column)
            v = c.value
            if hasattr(v, 'text') and not isinstance(v, str):     # ArrayFormula
                v = None
            d.value = v
            if c.has_style:
                d.font = copy(c.font)
                d.fill = copy(c.fill)
                d.border = copy(c.border)
                d.alignment = copy(c.alignment)
                d.number_format = c.number_format
                d.protection = copy(c.protection)
    for k, cd in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[k].width = cd.width
        dst_ws.column_dimensions[k].hidden = cd.hidden
    for k, rd in src_ws.row_dimensions.items():
        if rd.height:
            dst_ws.row_dimensions[k].height = rd.height
    for mr in src_ws.merged_cells.ranges:
        dst_ws.merge_cells(str(mr))
    dst_ws.sheet_view.zoomScale = src_ws.sheet_view.zoomScale
    dst_ws.sheet_format = copy(src_ws.sheet_format)
    dst_ws.page_setup.orientation = src_ws.page_setup.orientation
    dst_ws.page_setup.paperSize = src_ws.page_setup.paperSize
    for dv in src_ws.data_validations.dataValidation:
        f1 = dv.formula1
        if f1 and '青岛公司' in f1 and '东鹏公司' in f1:        # 所属公司：原来写死 4 家，改成跟【基础资料·供应商】走
            f1 = '=供应商列表'
        n = DataValidation(type=dv.type, formula1=f1, allow_blank=dv.allow_blank,
                           showErrorMessage=dv.showErrorMessage, errorStyle=dv.errorStyle)
        dst_ws.add_data_validation(n)
        n.add(str(dv.sqref))


def _fix_month_sheet(ws, blocks, total_row, sheet_label):
    """blocks: [(首行, 末行, 小计行)]；total_row: 总计行"""
    hf = lambda r: f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(G{r})),E{r}*G{r},"")'   # 数量或单价不是数字（如「2箱」）→ 空，不报错
    for (r0, r1, sub) in blocks:
        for r in range(r0, r1 + 1):
            ws[f'H{r}'] = hf(r)
        ws[f'H{sub}'] = f'=SUM(H{r0}:H{r1})'
    # 总计行往下一直到第 1000 行也铺上公式（表头写着数据行范围 A6:J1000，写在下面的也要算）
    for r in range(total_row + 1, BS_ROWS[1] + 1):
        ws[f'H{r}'] = hf(r)
    # 总计 = 所有「填了产品名称」的明细行：小计、总计行 C 列是空的，不会重复加；漏填所属公司的也不会掉
    t = total_row
    ws[f'H{t}'] = (f'=SUMIF(C{BS_ROWS[0]}:C{t - 1},"<>",H{BS_ROWS[0]}:H{t - 1})'
                   f'+SUMIF(C{t + 1}:C{BS_ROWS[1]},"<>",H{t + 1}:H{BS_ROWS[1]})')
    ws['B2'].number_format = 'yyyy"年"m"月"'
    # 有数量没单价：金额会悄悄变成 0，标红提醒；没填所属公司的也标
    ws.conditional_formatting.add(f'E{BS_ROWS[0]}:G{BS_ROWS[1]}',
                                  FormulaRule(formula=[f'AND($C{BS_ROWS[0]}<>"",$E{BS_ROWS[0]}<>"",$G{BS_ROWS[0]}="")'], fill=FILL_WARN))
    ws.conditional_formatting.add(f'B{BS_ROWS[0]}:B{BS_ROWS[1]}',
                                  FormulaRule(formula=[f'AND($C{BS_ROWS[0]}<>"",$B{BS_ROWS[0]}="")'], fill=FILL_WARN))


def build(wb, ctx, src2):
    # ── 搬四张表 ──
    for name in [SH_BSS, SH_BS6, SH_BS7, SH_BS8]:
        dst = wb.create_sheet(name)
        _copy_sheet(src2.wb[name], dst)
        dst.sheet_properties.tabColor = 'FF833C0C'[2:]

    # 报损明细台账（6 月）：H 列原来是 F×G（单位×单价 → #VALUE!），且小计被覆盖；恢复成和 7 月一样的结构
    ws6 = wb[SH_BS6]
    _fix_month_sheet(ws6, [(6, 10, 11), (12, 21, 22), (23, 33, 34)], 35, '6月')
    ws6['I2'] = '=H35'
    ws6['I2'].number_format = MONEY
    ws6.freeze_panes = 'A6'
    # 7 月：公式本来就对，统一一下
    ws7 = wb[SH_BS7]
    _fix_month_sheet(ws7, [(6, 10, 11), (12, 21, 22), (23, 33, 34)], 35, '7月')
    ws7['I2'] = '=H35'
    ws7['I2'].number_format = MONEY
    # 8 月：总计 H46 漏加了「其他公司」小计 H33
    ws8 = wb[SH_BS8]
    _fix_month_sheet(ws8, [(6, 10, 11), (12, 22, 23), (24, 32, 33), (34, 44, 45)], 46, '8月')
    ws8.freeze_panes = 'A2'
    # 9 月：照 8 月复制一张空白表
    ws9 = wb.create_sheet(SH_BS9)
    _copy_sheet(src2.wb[SH_BS8], ws9)
    ws9.sheet_properties.tabColor = 'FF833C0C'[2:]
    _fix_month_sheet(ws9, [(6, 10, 11), (12, 22, 23), (24, 32, 33), (34, 44, 45)], 46, '9月')
    for r in range(6, 47):
        for col in 'EIJ':
            ws9[f'{col}{r}'] = None
    ws9['B2'] = dt.datetime(ctx['year'], 9, 1)
    ws9['I2'] = '=H46'
    ws9.freeze_panes = 'A2'
    for w in (ws6, ws7, ws8, ws9):
        w['I2'].number_format = MONEY

    _build_staging(wb, ctx)
    _build_summary(wb, ctx)


def _build_staging(wb, ctx):
    ws = wb[SH_AUX]
    base = f'{SH_BASE}!'
    put(ws, f'{BSX_NAME}3', '报损中转站：每个商品每月报损（按报损表品名匹配）', F_NOTE, border=False)
    for g, t in [(BSX_AMT0, '金额'), (BSX_RAW0, '原始数量'), (BSX_BOX0, '按箱/件录的数量'), (BSX_PCS0, '折成件')]:
        put(ws, f'{CL(g)}3', t, F_NOTE, border=False)
        for m in range(1, 13):
            col = CL(g + m - 1)
            ws[f'{col}4'] = m
            ws[f'{col}5'] = (f'=IFERROR(INDEX({base}${M_SHEET}${BSM_R0}:${M_SHEET}${BSM_R1},'
                             f'MATCH({m},{base}${M_MONTH}${BSM_R0}:${M_MONTH}${BSM_R1},0))&"","")')
    ws[f'{BSX_NAME}5'] = '报损品名'
    ws[f'{BSX_PACK}5'] = '每件数量'
    ws[f'{BSX_SEL}5'] = '所选月份金额'
    ws[f'{BSX_KEY}5'] = '排名键'
    ws[f'{BSX_SELQ}5'] = '所选月份数量'

    def ind(colref, col_letter):
        return f'INDIRECT("\'"&{colref}&"\'!{col_letter}{BS_ROWS[0]}:{col_letter}{BS_ROWS[1]}")'

    def safe(expr):                 # 月表名写错 / 表里有怪值 → 当 0，由【数据校验】报出来，不让整本报表跟着报错
        return f'IFERROR({expr},0)'

    sel = f'{SH_BSS}!$A$3'
    for i in range(BASE_R1 - BASE_R0 + 1):
        r = AUX_BS_P0 + i
        br = BASE_R0 + i
        nm, pk = at(base_rng(G_BSNAME, BASE_R0, BASE_R1), i + 1), at(base_rng(G_PACK, BASE_R0, BASE_R1), i + 1)
        ws[f'{BSX_NAME}{r}'] = f'=IF({nm}="","",{nm}&"")'
        ws[f'{BSX_PACK}{r}'] = f'=IF(N({pk})<=0,1,{pk})'
        for m in range(1, 13):
            a, raw, box, pcs = (CL(BSX_AMT0 + m - 1), CL(BSX_RAW0 + m - 1), CL(BSX_BOX0 + m - 1), CL(BSX_PCS0 + m - 1))
            sh = f'{a}$5'
            skip = f'OR(${BSX_NAME}{r}="",{sh}="")'
            ws[f'{a}{r}'] = f'=IF({skip},0,' + safe(f'SUMIF({ind(sh, "C")},${BSX_NAME}{r},{ind(sh, "H")})') + ')'
            ws[f'{raw}{r}'] = f'=IF({skip},0,' + safe(f'SUMIF({ind(sh, "C")},${BSX_NAME}{r},{ind(sh, "E")})') + ')'
            ws[f'{box}{r}'] = (f'=IF({skip},0,' + safe(f'SUMIFS({ind(sh, "E")},{ind(sh, "C")},${BSX_NAME}{r},{ind(sh, "F")},"箱")'
                               f'+SUMIFS({ind(sh, "E")},{ind(sh, "C")},${BSX_NAME}{r},{ind(sh, "F")},"件")') + ')')
            ws[f'{pcs}{r}'] = f'=({raw}{r}-{box}{r})/${BSX_PACK}{r}+{box}{r}'
        amt_rng = f'{bsx_amt(1)}{r}:{bsx_amt(12)}{r}'
        raw_rng = f'{CL(BSX_RAW0)}{r}:{CL(BSX_RAW0 + 11)}{r}'
        mask = f'((${bsx_amt(1)}$4:${bsx_amt(12)}$4={sel})+({sel}="全年"))'
        ws[f'{BSX_SEL}{r}'] = f'=SUMPRODUCT({mask}*{amt_rng})'
        ws[f'{BSX_SELQ}{r}'] = f'=SUMPRODUCT({mask}*{raw_rng})'
        ws[f'{BSX_KEY}{r}'] = f'=IF({BSX_SEL}{r}>0,{BSX_SEL}{r}+(1000-ROW())/10000000,"")'
    # 供应商块
    put(ws, f'{BSX_NAME}{AUX_BS_S0 - 1}', '报损中转站：每家公司每月报损金额', F_NOTE, border=False)
    for i in range(SUP_R1 - SUP_R0 + 1):
        r = AUX_BS_S0 + i
        sr = SUP_R0 + i
        nm = at(SUP_NAMES, i + 1)
        ws[f'{BSX_NAME}{r}'] = f'=IF({nm}="","",{nm}&"")'
        for m in range(1, 13):
            a = bsx_amt(m)
            ws[f'{a}{r}'] = (f'=IF(OR(${BSX_NAME}{r}="",{a}$5=""),0,' +
                             safe(f'SUMIF({ind(f"{a}$5", "B")},${BSX_NAME}{r},{ind(f"{a}$5", "H")})') + ')')
        mask = f'((${bsx_amt(1)}$4:${bsx_amt(12)}$4={sel})+({sel}="全年"))'
        ws[f'{BSX_SEL}{r}'] = f'=SUMPRODUCT({mask}*{bsx_amt(1)}{r}:{bsx_amt(12)}{r})'
    # 每月总计 / 没对上商品档案的
    ws[f'{BSX_NAME}{AUX_BS_TOT}'] = '每月报损总金额'
    ws[f'{BSX_NAME}{AUX_BS_TOTQ}'] = '每月报损总数量'
    ws[f'{BSX_NAME}{AUX_BS_UNM}'] = '没对上商品档案的金额'
    ws[f'{BSX_NAME}{AUX_BS_NOPRICE}'] = '有数量没单价的行数'
    ws[f'{BSX_NAME}{AUX_BS_NOCOMP}'] = '有产品没填所属公司的行数'
    ws[f'{BSX_NAME}{AUX_BS_NONNUM}'] = '数量/单价不是数字的格数'
    for m in range(1, 13):
        a = bsx_amt(m)
        sh = f'{a}$5'
        ws[f'{a}{AUX_BS_TOT}'] = f'=IF({sh}="",0,' + safe(f'SUMIF({ind(sh, "C")},"<>",{ind(sh, "H")})') + ')'
        ws[f'{a}{AUX_BS_TOTQ}'] = f'=IF({sh}="",0,' + safe(f'SUMIF({ind(sh, "C")},"<>",{ind(sh, "E")})') + ')'
        ws[f'{a}{AUX_BS_UNM}'] = f'={a}{AUX_BS_TOT}-SUM({a}{AUX_BS_P0}:{a}{AUX_BS_P1})'
        ws[f'{a}{AUX_BS_NOPRICE}'] = f'=IF({sh}="",0,' + safe(f'COUNTIFS({ind(sh, "C")},"<>",{ind(sh, "E")},"<>",{ind(sh, "G")},"")') + ')'
        ws[f'{a}{AUX_BS_NOCOMP}'] = f'=IF({sh}="",0,' + safe(f'COUNTIFS({ind(sh, "C")},"<>",{ind(sh, "B")},"")') + ')'
        ws[f'{a}{AUX_BS_NONNUM}'] = f'=IF({sh}="",0,' + safe(f'SUMPRODUCT(ISTEXT({ind(sh, "E")})*1)+SUMPRODUCT(ISTEXT({ind(sh, "G")})*1)') + ')'
        ws[f'{a}{AUX_BS_NOH}'] = f'=IF({sh}="",0,' + safe(f'COUNTIFS({ind(sh, "E")},"<>",{ind(sh, "G")},"<>",{ind(sh, "H")},"")') + ')'
    mask = f'((${bsx_amt(1)}$4:${bsx_amt(12)}$4={sel})+({sel}="全年"))'
    for rr in (AUX_BS_TOT, AUX_BS_TOTQ, AUX_BS_UNM):
        ws[f'{BSX_SEL}{rr}'] = f'=SUMPRODUCT({mask}*{bsx_amt(1)}{rr}:{bsx_amt(12)}{rr})'


def _build_summary(wb, ctx):
    ws = wb[SH_BSS]
    aux = f'{SH_AUX}!'
    # 清掉原来的公式（含 UNIQUE/XLOOKUP 数组公式）
    for r in range(2, 42):
        for c in range(1, 6):
            v = ws.cell(r, c).value
            if (isinstance(v, str) and v.startswith('=')) or (v is not None and not isinstance(v, (str, int, float))):
                ws.cell(r, c).value = None
    ws['C37'] = None                         # 原来这里残留一个「艾精隆」
    # 统计月份选择（放在第 2/3 行最左边，原来空着）
    _st = copy(ws['B2']._style)
    ws['A2']._style = copy(_st)
    ws['A2'] = '统计月份'
    ws['A3']._style = copy(ws['B3']._style)
    ws['A3'].fill = FILL_SEL
    ws['A3'] = '全年'
    ws['A3'].number_format = '0"月";;;@'
    add_list_dv(ws, 'A3', '"全年,1,2,3,4,5,6,7,8,9,10,11,12"', '选「全年」或某个月')
    ws.column_dimensions['A'].width = max(ws.column_dimensions['A'].width or 8, 11)
    # KPI
    ws['B3'] = f'={aux}{BSX_SEL}{AUX_BS_TOT}'
    ws['C3'] = f'={aux}{BSX_SEL}{AUX_BS_TOTQ}'
    ws['D3'] = f'=COUNTIF({aux}{BSX_SEL}{AUX_BS_P0}:{BSX_SEL}{AUX_BS_P1},">0")'
    ws['E3'] = f'=COUNTIF({aux}{BSX_SEL}{AUX_BS_S0}:{BSX_SEL}{AUX_BS_S1},">0")+IF(ROUND(C12,2)>0,1,0)'
    ws['B3'].number_format = MONEY
    ws['C3'].number_format = '#,##0'
    # 分公司：第 7..11 行 = 供应商清单前 5 家，第 12 行 = 其余/没填或没登记的公司（保证合计 = 总报损金额）
    for i in range(5):
        r = 7 + i
        ws[f'A{r}'] = i + 1
        ws[f'B{r}'] = f'=IF({at(SUP_NAMES, i + 1)}="","",{at(SUP_NAMES, i + 1)})'
        ws[f'C{r}'] = f'=IF(B{r}="","",{aux}{BSX_SEL}{AUX_BS_S0 + i})'
        ws[f'D{r}'] = f'=IF(OR(B{r}="",N($C$13)=0),"",C{r}/$C$13)'
        ws[f'C{r}'].number_format = MONEY
        ws[f'D{r}'].number_format = '0.0%'
    ws['A12'] = 6
    ws['B12'] = '其他/未登记公司'
    ws['C12'] = '=B3-SUM(C7:C11)'
    ws['C12'].number_format = MONEY
    ws['D12'] = '=IF(N($C$13)=0,"",C12/$C$13)'
    ws['D12'].number_format = '0.0%'
    ws['C13'] = '=SUM(C7:C12)'
    ws['D13'] = '=IF(N(C13)=0,"",SUM(D7:D12))'
    ws['D13'].number_format = '0.0%'
    ws['A5'] = '分公司报损汇总（按所选月份）'
    # 分品项：第 17..36 行，按金额从大到小
    ws['A15'] = '分品项报损汇总（按金额降序 · 所选月份）'
    keys = f'{aux}${BSX_KEY}${AUX_BS_P0}:${BSX_KEY}${AUX_BS_P1}'
    for k in range(1, 21):
        r = 16 + k
        ws[f'A{r}'] = k
        ws[f'F{r}'] = f'=IFERROR(MATCH(LARGE({keys},{k}),{keys},0),"")'
        ws[f'B{r}'] = f'=IF($F{r}="","",INDEX({aux}${BSX_NAME}${AUX_BS_P0}:${BSX_NAME}${AUX_BS_P1},$F{r}))'
        ws[f'C{r}'] = f'=IF($F{r}="","",INDEX({SH_BASE}!${G_SUP}${BASE_R0}:${G_SUP}${BASE_R1},$F{r})&"")'
        ws[f'D{r}'] = f'=IF($F{r}="","",INDEX({aux}${BSX_SEL}${AUX_BS_P0}:${BSX_SEL}${AUX_BS_P1},$F{r}))'
        ws[f'E{r}'] = f'=IF(OR($F{r}="",N($D$41)=0),"",D{r}/$D$41)'
        ws[f'D{r}'].number_format = MONEY
        ws[f'E{r}'].number_format = '0.0%'
    ws.column_dimensions['F'].hidden = True
    _s = copy(ws['B36']._style)
    ws['B37']._style = copy(_s)
    ws['B37'] = '其余品项 / 没对上商品档案的'
    ws['D37']._style = copy(ws['D36']._style)
    ws['D37'] = '=D41-SUM(D17:D36)'
    ws['D37'].number_format = MONEY
    ws['D41'] = '=B3'
    ws['D41'].number_format = MONEY
    ws['E41'] = '=IF(N(D41)=0,"",1)'
    ws['E41'].number_format = '0%'
    ws['A1'] = '产品过期报损汇总一览'
    ws.freeze_panes = 'A4'
