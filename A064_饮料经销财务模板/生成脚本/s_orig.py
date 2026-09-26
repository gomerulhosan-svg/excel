# -*- coding: utf-8 -*-
"""原表 1 的几张：出库明细 / 存条明细 / 总览汇总 / 商品库存 / 客户查询 —— 样子不动，只补公式、修错；
外加新表【采购进货】和隐藏的【_辅助】（商品库存自动列组合、客户查询取数用）。"""
from copy import copy
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import PatternFill

from common import *


def _copy_style(src, dst):
    dst._style = copy(src._style)


def _copy_row_style(ws, src_row, dst_row, c0=1, c1=8):
    for c in range(c0, c1 + 1):
        _copy_style(ws.cell(src_row, c), ws.cell(dst_row, c))
    if ws.row_dimensions[src_row].height:
        ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


# ───────────────────────── 出库明细 / 存条明细 ─────────────────────────
def _retitle(ws, text, last_col, color, tip, h2=None):
    """第 1 行标题带 + 第 2 行提示，统一成 A064 新表的样子（原来各张表颜色、字号、底色都不一样）"""
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row <= 2:
            ws.unmerge_cells(str(mr))
    for r in (1, 2):
        for c in range(1, CI(last_col) + 1):
            ws.cell(r, c).value = None if (r, c) != (1, 1) else ws.cell(r, c).value
    title(ws, text, last_col, color, tip, h2=h2)
    ws.sheet_properties.tabColor = color[2:]


def _fix_ledger(ws, r0, r1, ctx, changes, sheet, color, heads, title_text, tip):
    """名称规范化、金额公式铺满、序号改公式、下拉、筛选；颜色格式统一：白底细框，公式列灰底"""
    widths(ws, {'A': 7, 'B': 12, 'C': 22, 'D': 14, 'E': 12, 'F': 12, 'G': 15, 'H': 22})
    _retitle(ws, title_text, 'H', color, tip)
    header(ws, 3, list(zip('ABCDEFGH', heads)), color, height=34)
    for r in range(r0, r1 + 1):
        # 名称规范化（fixes.py 里写明的）
        for col, mp, kind in [('C', ctx['cus_merge'], '客户'), ('D', ctx['goods_merge'], '商品')]:
            v = ws[f'{col}{r}'].value
            if isinstance(v, str):
                s = v.strip()
                new = mp.get(s, s)
                if new != v:
                    changes.append((sheet, f'{col}{r}', kind, v, new))
                    ws[f'{col}{r}'].value = new
        # 序号：原来手打的（你上一版把重号、跳号重新排成了 1、2、3…），现在改成公式，插行删行也不会乱
        ws[f'A{r}'].value = f'=IF(AND(B{r}="",C{r}="",D{r}=""),"",ROW()-{r0 - 1})'
        g = ws[f'G{r}'].value
        if g is None or (isinstance(g, str) and g.startswith('=')):
            ws[f'G{r}'].value = f'=IF(AND(E{r}<>"",F{r}<>""),E{r}*F{r},"")'
    style_rows(ws, r0, r1, 'ABCDEFGH', auto=('A', 'G'),
               fmts={'B': DATE, 'E': QTY, 'F': PRICE, 'G': MONEY}, aligns={'H': AL, 'G': AR_}, height=16)
    add_date_dv(ws, f'B{r0}:B{r1}')
    add_list_dv(ws, f'C{r0}:C{r1}', '=客户列表', '从下拉选客户；新客户先去【总览汇总】B 列加')
    add_list_dv(ws, f'D{r0}:D{r1}', '=商品列表', '从下拉选商品；新商品先去【基础资料·商品档案】加')
    ws.auto_filter.ref = f'A3:H{r1}'
    ws.freeze_panes = 'A4'


def fix_out(wb, ctx, changes):
    ws = wb[SH_OUT]
    _fix_ledger(ws, OUT_R0, OUT_R1, ctx, changes, SH_OUT, C_OUT,
                ['序号', '日期', '客户名称', '商品品类', '领用数量(件)', '单价(元/件)', '本次领用金额(元)', '备注'],
                '商品领用明细台账',
                '💡 每次送货/拣货在此登记（客户、商品从下拉选），总览表、库存表、利润表、对账单自动更新')
    import fixes
    for r, d in fixes.OUT_DATE_FIX.items():
        old = ws[f'B{r}'].value
        if isinstance(old, str):
            ws[f'B{r}'].value = d
            changes.append((SH_OUT, f'B{r}', '日期', old, d.strftime('%Y-%m-%d')))


def fix_cun(wb, ctx, changes):
    ws = wb[SH_CUN]
    _fix_ledger(ws, CUN_R0, CUN_R1, ctx, changes, SH_CUN, C_CUN,
                ['序号', '入库日期', '客户名称', '商品品类', '入库数量(件)', '入库单价(元/件)', '本次入库总金额(元)', '备注'],
                '商品入库明细台账（客户存条）',
                '💡 每次客户存条在此登记，一个商品一行，库存表按「客户+商品」自动汇总累加')


# ───────────────────────── 采购进货（新） ─────────────────────────
def build_buy(wb, ctx):
    ws = wb.create_sheet(SH_BUY)
    ws.sheet_properties.tabColor = C_BUY[2:]
    title(ws, '采购进货明细台账', 'J', C_BUY,
          '💡 从厂家/供应商进货在这里登记，一个商品一行：数量×进价自动出金额；厂家搭赠的货填「赠品数量」，只加库存不加货款；'
          '退货数量填负数。付货款去【资金台帐】记一笔并选供应商。公司库存、成本、利润表、供应商对账单都从这里取数。')
    header(ws, BUY_HDR, [(B_SEQ, '序号'), (B_DATE, '日期'), (B_SUP, '供应商'), (B_GOODS, '商品品类'),
                         (B_QTY, '数量(件)'), (B_PRICE, '进价(元/件)'), (B_AMT, '金额(元)'),
                         (B_GIFT, '赠品数量(件)'), (B_NOTE, '备注'), (B_CHK, '校验')], C_BUY)
    rows = ctx.get('buy_rows', [])
    for i in range(BUY_R1 - BUY_R0 + 1):
        r = BUY_R0 + i
        d = rows[i] if i < len(rows) else {}
        put(ws, f'{B_SEQ}{r}', f'=IF(AND({B_SUP}{r}="",{B_GOODS}{r}=""),"",ROW()-{BUY_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{B_DATE}{r}', d.get('date'), F_TXT, fmt=DATE, align=AC)
        put(ws, f'{B_SUP}{r}', d.get('supplier'), F_TXT, align=AC)
        put(ws, f'{B_GOODS}{r}', d.get('goods'), F_TXT, align=AC)
        put(ws, f'{B_QTY}{r}', d.get('qty'), F_TXT, fmt=QTY, align=AC)
        put(ws, f'{B_PRICE}{r}', d.get('price'), F_TXT, fmt=PRICE, align=AC)
        put(ws, f'{B_AMT}{r}', f'=IF(AND({B_QTY}{r}<>"",{B_PRICE}{r}<>""),{B_QTY}{r}*{B_PRICE}{r},"")', F_AUTO, FILL_AUTO, MONEY, AC)
        put(ws, f'{B_GIFT}{r}', d.get('gift'), F_TXT, fmt=QTY, align=AC)
        put(ws, f'{B_NOTE}{r}', d.get('note'), F_TXT, align=AL)
        blank = f'AND({B_DATE}{r}="",{B_SUP}{r}="",{B_GOODS}{r}="",{B_QTY}{r}="",{B_PRICE}{r}="",{B_GIFT}{r}="")'
        put(ws, f'{B_CHK}{r}',
            f'=IF({blank},"",IF(NOT(ISNUMBER({B_DATE}{r})),"✗ 日期不是真日期",IF(OR({B_DATE}{r}<年初日,{B_DATE}{r}>=年末日+1),"✗ 日期不在本年度",'
            f'IF({B_SUP}{r}="","✗ 没选供应商",IF(COUNTIF({SUP_NAMES},{B_SUP}{r})=0,"✗ 供应商不在基础资料里",'
            f'IF({B_GOODS}{r}="","✗ 没选商品",IF(COUNTIF({GOODS_NAMES},{B_GOODS}{r})=0,"✗ 商品不在商品档案里",'
            f'IF(AND({B_QTY}{r}="",{B_GIFT}{r}=""),"✗ 没填数量",IF(AND(N({B_QTY}{r})<>0,{B_PRICE}{r}=""),"✗ 没填进价","√")))))))))',
            F_AUTO, FILL_AUTO, align=AL)
    ws.conditional_formatting.add(f'{B_CHK}{BUY_R0}:{B_CHK}{BUY_R1}',
                                  FormulaRule(formula=[f'LEFT(${B_CHK}{BUY_R0},1)="✗"'], fill=FILL_WARN))
    add_date_dv(ws, f'{B_DATE}{BUY_R0}:{B_DATE}{BUY_R1}')
    add_list_dv(ws, f'{B_SUP}{BUY_R0}:{B_SUP}{BUY_R1}', '=供应商列表', '从下拉选供应商')
    add_list_dv(ws, f'{B_GOODS}{BUY_R0}:{B_GOODS}{BUY_R1}', '=商品列表', '从下拉选商品')
    widths(ws, {B_SEQ: 7, B_DATE: 13, B_SUP: 15, B_GOODS: 15, B_QTY: 11, B_PRICE: 12, B_AMT: 15,
                B_GIFT: 12, B_NOTE: 22, B_CHK: 22})
    ws.auto_filter.ref = f'A{BUY_HDR}:{B_CHK}{BUY_R1}'
    ws.freeze_panes = f'A{BUY_R0}'
    return ws


# ───────────────────────── 总览汇总（客户名单） ─────────────────────────
def fix_ov(wb, ctx):
    ws = wb[SH_OV]
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= 3:
            ws.unmerge_cells(str(mr))
    for r in range(3, max(ws.max_row, OV_TOT) + 1):
        for c in range(1, 21):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy(ws.cell(OV_TOT + 30, 20)._style)
    widths(ws, {'A': 6, 'B': 20, 'C': 15, 'D': 15, 'E': 15, 'F': 15, 'G': 15, 'H': 9, 'I': 14, 'J': 15, 'K': 14})
    _retitle(ws, '客户存条&货款总览汇总表', 'K', C_OV,
             '💡 一篮子查看所有客户情况：新增客户在 B 列往下加（C 填初始存条，K 填年初就欠的货款），领用和收款自动更新。'
             '第 4 行是合计（往下翻也看得到）。')
    header(ws, 3, list(zip('ABCDEFGHIJK', ['序号', '客户名称', '初始存条金额(元)', '已领用金额(元)', '剩余存条金额(元)',
                                            '已付款金额(元)', '未付货款(元)', '状态', '库存总结余数量(件)',
                                            '库存总结余金额(元)', '期初欠款(元)'])), C_OV, height=40)

    cust = ctx['customers']       # [(name, deposit, src_rows)]
    for i in range(OV_R1 - OV_R0 + 1):
        r = OV_R0 + i
        name, dep = (cust[i][0], cust[i][1]) if i < len(cust) else (None, None)
        ws[f'A{r}'] = f'=IF(B{r}="","",COUNTA($B${OV_R0}:B{r}))'
        ws[f'B{r}'] = name
        ws[f'C{r}'] = dep
        ws[f'D{r}'] = f'=IF(B{r}<>"",SUMIF({SH_OUT}!C:C,B{r},{SH_OUT}!G:G),"")'
        ws[f'E{r}'] = f'=IF(AND(C{r}<>"",D{r}<>""),C{r}-D{r},"")'
        ws[f'F{r}'] = (f'=IF(B{r}<>"",SUMIFS({SH_CASH}!${K_NET}:${K_NET},{SH_CASH}!${K_CUS}:${K_CUS},B{r},'
                       f'{SH_CASH}!${K_TO}:${K_TO},"{TO_AR}"),"")')
        ws[f'G{r}'] = f'=IF(B{r}="","",N(K{r})+N(D{r})-N(F{r}))'
        ws[f'H{r}'] = f'=IF(B{r}="","",IF(ROUND(G{r},2)=0,"已结清",IF(G{r}>0,"有欠款","预存款")))'
        ws[f'I{r}'] = f'=IF(B{r}="","",IFERROR(SUMIF({SH_INV}!B:B,B{r},{SH_INV}!I:I),0))'
        ws[f'J{r}'] = f'=IF(B{r}="","",IFERROR(SUMIF({SH_INV}!B:B,B{r},{SH_INV}!J:J),0))'
    money = {c: MONEY for c in 'CDEFGJK'}
    style_rows(ws, OV_R0, OV_R1, 'ABCDEFGHIJK', auto=tuple('ADEFGHIJ'), fmts={**money, 'I': QTY},
               aligns={c: AR_ for c in 'CDEFGIJK'}, bold=('G',), height=16)
    # 合计：底部一行 + 顶上第 4 行（冻结在表头下面）
    for rr, lab in ((OV_TOT, '合计'), (OV_TOP, '合计')):
        for c in range(1, 12):
            col = CL(c)
            put(ws, f'{col}{rr}', None, F_TXTB, FILL_TOT, money.get(col, QTY if col == 'I' else None), AR_)
        ws[f'B{rr}'] = lab
        ws[f'B{rr}'].alignment = AC
    for col in 'CDEFGIJK':
        ws[f'{col}{OV_TOT}'] = f'=SUM({col}{OV_R0}:{col}{OV_R1})'
        ws[f'{col}{OV_TOP}'] = f'={col}{OV_TOT}'
    ws.row_dimensions[OV_TOP].height = 22
    # 重名提醒：同一个客户写了两行，SUMIF 会把他的领用/收款算两遍
    ws.conditional_formatting.add(f'B{OV_R0}:B{OV_R1}',
                                  FormulaRule(formula=[f'AND(B{OV_R0}<>"",COUNTIF($B${OV_R0}:$B${OV_R1},B{OV_R0})>1)'],
                                              fill=FILL_WARN))
    ws.conditional_formatting.add(f'H{OV_R0}:H{OV_R1}', FormulaRule(formula=[f'$H{OV_R0}="有欠款"'],
                                                                   font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.conditional_formatting.add(f'H{OV_R0}:H{OV_R1}', FormulaRule(formula=[f'$H{OV_R0}="预存款"'],
                                                                   font=Font(name=YH, sz=10, color='FF00B050')))
    ws.auto_filter.ref = f'A{OV_TOP}:K{OV_R1}'
    ws.freeze_panes = 'A5'


# ───────────────────────── _辅助 + 商品库存（自动列出组合） ─────────────────────────
AUX_COMBO_R0, AUX_COMBO_R1 = 4, 4003      # A 列：存条 1000 行 + 出库 3000 行
AUX_Q_COL = 'D'                           # 客户查询取数


def build_aux_combo(wb, ctx):
    ws = wb[SH_AUX]
    put(ws, 'A3', '商品库存组合键', F_NOTE, border=False)
    put(ws, 'B3', '第k个组合在A列的位置', F_NOTE, border=False)
    put(ws, 'D3', '客户查询：第几个', F_NOTE, border=False)
    put(ws, 'A2', f'=COUNT($A${AUX_COMBO_R0}:$A${AUX_COMBO_R1})', F_NOTE, border=False)   # 组合个数
    cus = f'{SH_OV}!$B${OV_R0}:$B${OV_R1}'
    # 注意：这里全部「按位置」取明细表的第 k 行（INDEX），不直接写 存条明细!C100 这种地址——
    # 用户在明细表里删一行/插一行，区域会跟着伸缩，位置始终对得上，不会出 #REF! 也不会串行
    for sh, r0, r1, a0 in ((SH_CUN, CUN_R0, CUN_R1, AUX_COMBO_R0), (SH_OUT, OUT_R0, OUT_R1, AUX_COMBO_R0 + CUN_R1 - CUN_R0 + 1)):
        C = f'{sh}!$C${r0}:$C${r1}'
        Dg = f'{sh}!$D${r0}:$D${r1}'
        for k in range(1, r1 - r0 + 2):
            r = a0 + k - 1
            c, d = at(C, k), at(Dg, k)
            key = (f'IFERROR(MATCH({c},{cus},0),900)*10000000+IFERROR(MATCH({d},{GOODS_NAMES},0),900)*10000+ROW()')
            dup = f'COUNTIFS({upto(C, k)},{c},{upto(Dg, k)},{d})>1'
            if sh == SH_OUT:        # 出库里的组合，存条里已经有的就不重复列
                dup = f'OR(COUNTIFS({cun("C")},{c},{cun("D")},{d})>0,{dup})'
            ws[f'A{r}'] = f'=IF(OR({c}="",{d}=""),"",IF({dup},"",{key}))'
    for k in range(1, INV_R1 - INV_R0 + 2):
        r = AUX_COMBO_R0 + k - 1
        ws[f'B{r}'] = (f'=IF({k}>$A$2,"",IFERROR(MATCH(SMALL($A${AUX_COMBO_R0}:$A${AUX_COMBO_R1},{k}),'
                       f'$A${AUX_COMBO_R0}:$A${AUX_COMBO_R1},0),""))')
    # 公司库存「客户存条未提」：商品库存第 k 行这个组合，截至【公司库存】所选月底还没提走的数量（多提的按 0，不能抵别人的）
    put(ws, 'R3', '存条未提(公司库存月底)', F_NOTE, border=False)
    end = f'DATE(年度,{SH_STK}!$C$3+1,1)'
    IB, IC = f'{SH_INV}!$B${INV_R0}:$B${INV_R1}', f'{SH_INV}!$C${INV_R0}:$C${INV_R1}'
    for k in range(1, INV_R1 - INV_R0 + 2):
        r = AUX_COMBO_R0 + k - 1
        cb, cc = at(IB, k), at(IC, k)
        cond = lambda sh: f'{sh}!$C:$C,{cb},{sh}!$D:$D,{cc},{sh}!$B:$B,"<"&{end}'
        ws[f'R{r}'] = (f'=IF({cb}="",0,MAX(0,SUMIFS({SH_CUN}!$E:$E,{cond(SH_CUN)})'
                       f'-SUMIFS({SH_OUT}!$E:$E,{cond(SH_OUT)})))')
        # 客户查询：商品库存第 k 行是不是查询的客户，是的话是第几个
        ws[f'{AUX_Q_COL}{r}'] = (f'=IF(AND({cb}<>"",{cb}={SH_Q}!$B$4),COUNTIF({upto(IB, k)},{SH_Q}!$B$4),"")')


def fix_inv(wb, ctx):
    ws = wb[SH_INV]
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row >= 3:
            ws.unmerge_cells(str(mr))
    blank = copy(ws.cell(INV_TOT + 30, 20)._style)
    for r in range(3, max(ws.max_row, INV_TOT) + 1):
        for c in range(1, 21):
            ws.cell(r, c).value = None
            ws.cell(r, c)._style = copy(blank)
    widths(ws, {'A': 6, 'B': 20, 'C': 14, 'D': 12, 'E': 11, 'F': 14, 'G': 12, 'H': 14, 'I': 12, 'J': 14, 'K': 12})
    _retitle(ws, '商品库存明细（按客户+商品）', 'K', C_INV,
             '💡 每个客户每个商品的存条库存：全自动，存条/出库里出现新的「客户+商品」这里自动多一行；剩余数量为负（领的比存的多）标红')
    header(ws, INV_HDR, list(zip('ABCDEFGHIJK', ['序号', '客户名称', '商品品类', '累计入库数量(件)', '入库单价(元)',
                                                 '累计入库总金额(元)', '已领用数量(件)', '已领用金额(元)', '剩余数量(件)',
                                                 '剩余金额(元)', '最后更新日期'])), C_INV, height=40)
    pos = lambda r: f'{SH_AUX}!$B{r}'
    for r in range(INV_R0, INV_R1 + 1):
        p = pos(r)
        ws[f'A{r}'] = f'=IF(B{r}="","",ROW()-{INV_R0 - 1})'
        ws[f'B{r}'] = (f'=IF({p}="","",IF({p}<={CUN_R1 - CUN_R0 + 1},INDEX({cun("C")},{p}),'
                       f'INDEX({out("C")},{p}-{CUN_R1 - CUN_R0 + 1})))')
        ws[f'C{r}'] = (f'=IF({p}="","",IF({p}<={CUN_R1 - CUN_R0 + 1},INDEX({cun("D")},{p}),'
                       f'INDEX({out("D")},{p}-{CUN_R1 - CUN_R0 + 1})))')
        cond = lambda sh: f'{sh}!C:C,$B{r},{sh}!D:D,$C{r}'
        ws[f'D{r}'] = f'=IF($B{r}="","",IFERROR(SUMIFS({SH_CUN}!E:E,{cond(SH_CUN)}),0))'
        ws[f'E{r}'] = (f'=IF(OR($B{r}="",N(D{r})=0),"",IFERROR(ROUND(SUMIFS({SH_CUN}!G:G,{cond(SH_CUN)})'
                       f'/SUMIFS({SH_CUN}!E:E,{cond(SH_CUN)},{SH_CUN}!F:F,">0"),2),""))')
        ws[f'F{r}'] = f'=IF($B{r}="","",IFERROR(SUMIFS({SH_CUN}!G:G,{cond(SH_CUN)}),0))'
        ws[f'G{r}'] = f'=IF($B{r}="","",IFERROR(SUMIFS({SH_OUT}!E:E,{cond(SH_OUT)}),0))'
        ws[f'H{r}'] = f'=IF($B{r}="","",IFERROR(SUMIFS({SH_OUT}!G:G,{cond(SH_OUT)}),0))'
        ws[f'I{r}'] = f'=IF($B{r}="","",D{r}-G{r})'
        ws[f'J{r}'] = f'=IF($B{r}="","",F{r}-H{r})'
        agg = lambda sh, r0, r1: (f'IFERROR(_xlfn.AGGREGATE(14,6,{sh}!$B${r0}:$B${r1}/(({sh}!$C${r0}:$C${r1}=$B{r})'
                                  f'*({sh}!$D${r0}:$D${r1}=$C{r})),1),0)')     # 最后一次进/出的日期（AGGREGATE：Excel 2010 起都有）
        ws[f'K{r}'] = f'=IF($B{r}="","",MAX({agg(SH_CUN, CUN_R0, CUN_R1)},{agg(SH_OUT, OUT_R0, OUT_R1)}))'
    style_rows(ws, INV_R0, INV_R1, 'ABCDEFGHIJK',
               fmts={'D': QTY, 'E': MONEY, 'F': MONEY, 'G': QTY, 'H': MONEY, 'I': QTY, 'J': MONEY, 'K': 'yyyy/mm/dd;;'},
               aligns={c: AR_ for c in 'DEFGHIJ'}, bold=('I', 'J'), height=16)
    r = INV_TOT
    for c in range(1, 12):
        put(ws, f'{CL(c)}{r}', None, F_TXTB, FILL_TOT, align=AR_)
    ws.merge_cells(f'A{r}:C{r}')
    ws[f'A{r}'] = '合计'
    ws[f'A{r}'].alignment = AC
    for col in 'DFGHIJ':
        ws[f'{col}{r}'] = f'=SUM({col}{INV_R0}:{col}{INV_R1})'
        ws[f'{col}{r}'].number_format = MONEY if col in 'FHJ' else QTY
    # 剩余数量为负：领用比存条多，标红
    ws.conditional_formatting.add(f'I{INV_R0}:J{INV_R1}',
                                  FormulaRule(formula=[f'AND(ISNUMBER($I{INV_R0}),$I{INV_R0}<0)'],
                                              font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.auto_filter.ref = f'A{INV_HDR}:K{INV_R1}'
    ws.freeze_panes = 'A4'


# ───────────────────────── 客户查询（原表，修好） ─────────────────────────
Q_R0, Q_R1 = 12, 51                      # 一个客户最多列 40 个商品


def _bar(ws, row, text, color, c1='G'):
    for mr in list(ws.merged_cells.ranges):
        if mr.min_row == row:
            ws.unmerge_cells(str(mr))
    ws.merge_cells(f'A{row}:{c1}{row}')
    put(ws, f'A{row}', text, Font(name=YH, sz=12, bold=True, color='FFFFFFFF'), fill(color), align=AL, border=False)
    ws.row_dimensions[row].height = 26


def fix_q(wb, ctx):
    ws = wb[SH_Q]
    for r in range(Q_R0, max(ws.max_row, Q_R1) + 1):
        for c in range(1, 10):
            ws.cell(r, c).value = None
    widths(ws, {'A': 20, 'B': 13, 'C': 15, 'D': 13, 'E': 15, 'F': 13, 'G': 15})
    _retitle(ws, '客户快速查询', 'G', C_Q,
             '💡 在黄色格子选客户名称，自动显示该客户的所有库存和货款情况（要打印给客户签字的对账单去【客户对账单】）')
    ws.row_dimensions[4].height = 28
    put(ws, 'A4', '🔍 请选客户名称：', F_KPI_L, align=AR_, border=False)
    for c in 'BCD':
        put(ws, f'{c}4', None, Font(name=YH, sz=12, bold=True, color='FF1F4E79'), FILL_SEL, align=AC)
    ws['B4'].value = ws['B4'].value or ctx['customers'][0][0]
    add_list_dv(ws, 'B4', '=客户列表', '选一个客户')
    put(ws, 'E4', '← 下拉选客户', F_NOTE, border=False)
    _bar(ws, 6, '💰 货款汇总', C_RP)
    header(ws, 7, list(zip('ABCDEFG', ['初始存条金额', '已领用金额', '剩余存条金额', '已付款金额', '未付货款', '状态', ''])),
           C_RP, height=24)
    for i, col in enumerate('ABCDEF'):
        v = f'VLOOKUP($B$4,{SH_OV}!$B${OV_R0}:$H${OV_R1},{i + 2},FALSE)'
        ws[f'{col}8'] = f'=IFERROR(IF({v}="","",{v}),"")'
    style_rows(ws, 8, 8, 'ABCDEFG', fmts={c: MONEY for c in 'ABCDE'}, bold=tuple('ABCDEF'), height=22)
    _bar(ws, 10, '📦 商品库存明细', C_INV)
    header(ws, 11, list(zip('ABCDEFG', ['商品品类', '初始数量', '初始金额', '已领用数量', '已领用金额', '剩余数量', '剩余金额'])),
           C_INV, height=24)
    src_cols = ['C', 'D', 'F', 'G', 'H', 'I', 'J']   # 商品 / 入库数 / 入库额 / 领用数 / 领用额 / 剩余数 / 剩余额
    for r in range(Q_R0, Q_R1 + 1):
        k = r - Q_R0 + 1
        ws[f'I{r}'] = f'=IFERROR(MATCH({k},{SH_AUX}!${AUX_Q_COL}${AUX_COMBO_R0}:${AUX_Q_COL}${AUX_COMBO_R0 + INV_R1 - INV_R0},0),"")'
        for j, sc in enumerate(src_cols):
            col = 'ABCDEFG'[j]
            ws[f'{col}{r}'] = f'=IF($I{r}="","",INDEX({SH_INV}!${sc}${INV_R0}:${sc}${INV_R1},$I{r}))'
    style_rows(ws, Q_R0, Q_R1, 'ABCDEFG', fmts={'B': QTY, 'C': MONEY, 'D': QTY, 'E': MONEY, 'F': QTY, 'G': MONEY},
               aligns={c: AR_ for c in 'BCDEFG'}, bold=('F', 'G'), height=16)
    ws.conditional_formatting.add(f'F{Q_R0}:G{Q_R1}', FormulaRule(formula=[f'AND(ISNUMBER($F{Q_R0}),$F{Q_R0}<0)'],
                                                                 font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    ws.column_dimensions['I'].hidden = True
