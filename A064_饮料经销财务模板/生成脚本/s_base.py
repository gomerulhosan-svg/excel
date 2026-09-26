# -*- coding: utf-8 -*-
"""【基础资料】：参数 / 商品档案 / 供应商 / 资金账户 / 收支项目 / 费用项目 / 报损月表清单，
外加整本册子要用的定义名称（年度、各下拉清单）。"""
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.formatting.rule import FormulaRule

from common import *
import fixes


def _block_title(ws, c0, c1, text, color):
    ws.merge_cells(f'{c0}3:{c1}3')
    put(ws, f'{c0}3', text, F_HDR, fill(color), align=AC)


def build(wb, ctx):
    ws = wb.create_sheet(SH_BASE)
    ws.sheet_properties.tabColor = C_BASE[2:]
    title(ws, '基础资料（参数 · 商品 · 供应商 · 资金账户 · 收支项目 · 费用项目 · 报损月表）', 'AO', C_BASE,
          '💡 淡黄格子手填，灰格子自动。这里的清单就是各录入表的下拉内容——清单里没有的先在这里加一行，再去录。'
          '客户名单还在【总览汇总】B 列（跟原来一样，新客户在那里加）。不要插入或删除行，直接往下面的空行里填。')
    ws.row_dimensions[3].height = 24
    ws.row_dimensions[4].height = 34

    # ── 参数 A:B ──
    _block_title(ws, 'A', 'B', '参数', C_BASE)
    header(ws, 4, [('A', '项目'), ('B', '值')], 'FF808080', height=34)
    params = [
        ('会计年度', fixes.BOOK_YEAR, True, '0'),
        ('公司名称', fixes.COMPANY_NAME, True, None),
        ('年初日', '=DATE(B5,1,1)', False, DATE),
        ('年末日', '=DATE(B5,12,31)', False, DATE),
        ('年初借款余额', 0, True, MONEY2),
        ('年初其他应收(+)/应付(-)', 0, True, MONEY2),
    ]
    for i, (k, v, is_in, fmt) in enumerate(params):
        r = 5 + i
        put(ws, f'A{r}', k, F_TXTB, FILL_SUBH, align=AL)
        put(ws, f'B{r}', v, F_IN if is_in else F_AUTO, FILL_IN if is_in else FILL_AUTO, fmt, AC)
    # 去向说明
    put(ws, 'A12', '「去向」说明（收支项目落到报表哪一行）', F_TXTB, FILL_SUBH, align=AL)
    ws.merge_cells('A12:B12')
    notes = [
        (TO_AR, '冲减客户未付货款（总览汇总·已付款）'),
        (TO_AP, '冲减欠供应商的货款'),
        (TO_EXP, '按费用项目进利润表'),
        (TO_REB, '利润表·其他业务收入'),
        (TO_OI, '利润表·营业外收入'),
        (TO_INV, '资产负债表·股东投入'),
        (TO_DRAW, '资产负债表·股东提取'),
        (TO_LOAN, '资产负债表·短期借款'),
        (TO_XFER, '账户互转，不进报表（收支要成对）'),
        (TO_OTH, '其他应收 / 其他应付（押金等）'),
    ]
    for i, (k, v) in enumerate(notes):
        r = 13 + i
        put(ws, f'A{r}', k, F_TXTB, align=AC)
        put(ws, f'B{r}', v, F_NOTE, align=ALW)
        ws.row_dimensions[r].height = 26

    # ── 商品档案 D:N ──
    _block_title(ws, G_SEQ, G_A0, '商品档案（出库 / 存条 / 采购 的「商品品类」都从这里选）', 'FF4472C4')
    header(ws, 4, [(G_SEQ, '序号'), (G_NAME, '商品名称'), (G_SUP, '供应商\n(所属公司)'),
                   (G_BSNAME, '报损表品名'), (G_SPEC, '规格'), (G_PACK, '每件\n瓶数'),
                   (G_COST, '参考进价\n(元/件)'), (G_PRICE, '参考售价\n(元/件)'),
                   (G_Q0, '期初库存\n(件)'), (G_P0, '期初单价\n(元/件)'), (G_A0, '期初金额\n(元)')], 'FF4472C4', height=34)
    goods = ctx['goods']
    for i in range(BASE_R1 - BASE_R0 + 1):
        r = BASE_R0 + i
        g = goods[i] if i < len(goods) else {}
        put(ws, f'{G_SEQ}{r}', f'=IF({G_NAME}{r}="","",ROW()-{BASE_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        for col, key, fmt in [(G_NAME, 'name', None), (G_SUP, 'supplier', None), (G_BSNAME, 'bs_name', None),
                              (G_SPEC, 'spec', None), (G_PACK, 'pack', '0'), (G_COST, 'cost', PRICE),
                              (G_PRICE, 'price', PRICE), (G_Q0, 'q0', QTY), (G_P0, 'p0', PRICE)]:
            v = g.get(key)
            put(ws, f'{col}{r}', v if v not in ('', None) else None, F_IN, FILL_IN, fmt, AC)
        put(ws, f'{G_A0}{r}', f'=IF(N({G_Q0}{r})=0,"",{G_Q0}{r}*N({G_P0}{r}))', F_AUTO, FILL_AUTO, MONEY2, AC)
    # 有商品名但没参考进价 → 标红提醒（成本会按 0 算）
    ws.conditional_formatting.add(f'{G_COST}{BASE_R0}:{G_COST}{BASE_R1}',
                                  FormulaRule(formula=[f'AND(${G_NAME}{BASE_R0}<>"",${G_COST}{BASE_R0}="")'], fill=FILL_WARN))
    add_list_dv(ws, f'{G_SUP}{BASE_R0}:{G_SUP}{BASE_R1}', '=供应商列表')

    # ── 供应商 P:T ──
    _block_title(ws, S_SEQ, S_NOTE, '供应商档案', 'FF7030A0')
    header(ws, 4, [(S_SEQ, '序号'), (S_NAME, '供应商名称'), (S_AP0, '年初应付(元)\n正=我欠他 负=预付'),
                   (S_CONTACT, '联系人/电话'), (S_NOTE, '备注')], 'FF7030A0', height=34)
    sups = ctx['suppliers']
    for i in range(SUP_R1 - SUP_R0 + 1):
        r = SUP_R0 + i
        s = sups[i] if i < len(sups) else {}
        put(ws, f'{S_SEQ}{r}', f'=IF({S_NAME}{r}="","",ROW()-{SUP_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{S_NAME}{r}', s.get('name'), F_IN, FILL_IN, align=AC)
        put(ws, f'{S_AP0}{r}', s.get('ap0'), F_IN, FILL_IN, MONEY2, AC)
        put(ws, f'{S_CONTACT}{r}', s.get('contact'), F_IN, FILL_IN, align=AC)
        put(ws, f'{S_NOTE}{r}', s.get('note'), F_IN, FILL_IN, align=ALW)

    # ── 资金账户 V:Z ──
    _block_title(ws, A_SEQ, A_NOTE, '资金账户（资金台帐「账户」下拉）', C_CASH)
    header(ws, 4, [(A_SEQ, '序号'), (A_NAME, '账户名称'), (A_BAL0, '年初余额(元)'),
                   (A_BALNOW, '当前余额(元)\n自动'), (A_NOTE, '备注')], C_CASH, height=34)
    accs = ctx['accounts']
    for i in range(ACC_R1 - ACC_R0 + 1):
        r = ACC_R0 + i
        a = accs[i] if i < len(accs) else {}
        put(ws, f'{A_SEQ}{r}', f'=IF({A_NAME}{r}="","",ROW()-{ACC_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{A_NAME}{r}', a.get('name'), F_IN, FILL_IN, align=AC)
        put(ws, f'{A_BAL0}{r}', a.get('bal0'), F_IN, FILL_IN, MONEY2, AC)
        put(ws, f'{A_BALNOW}{r}',
            f'=IF({A_NAME}{r}="","",N({A_BAL0}{r})+SUMIF({cash(K_ACC)},{A_NAME}{r},{cash(K_IN)})'
            f'-SUMIF({cash(K_ACC)},{A_NAME}{r},{cash(K_OUT)}))', F_AUTO, FILL_AUTO, MONEY2, AC)
        put(ws, f'{A_NOTE}{r}', a.get('note'), F_IN, FILL_IN, align=ALW)
    r = ACC_R1 + 1
    put(ws, f'{A_NAME}{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'{A_BAL0}{r}', f'=SUM({A_BAL0}{ACC_R0}:{A_BAL0}{ACC_R1})', F_TXTB, FILL_TOT, MONEY2, AC)
    put(ws, f'{A_BALNOW}{r}', f'=SUM({A_BALNOW}{ACC_R0}:{A_BALNOW}{ACC_R1})', F_TXTB, FILL_TOT, MONEY2, AC)

    # ── 收支项目 AB:AF ──
    _block_title(ws, C_SEQ, C_NOTE, '收支项目（资金台帐「收支项目」下拉；前 6 个名称别改，公式按名字认）', 'FF00B050')
    header(ws, 4, [(C_SEQ, '序号'), (C_NAME, '收支项目'), (C_DIR, '方向'), (C_TO, '去向'),
                   (C_NOTE, '说明')], 'FF00B050', height=34)
    cats = ctx['categories']
    for i in range(CAT_R1 - CAT_R0 + 1):
        r = CAT_R0 + i
        c = cats[i] if i < len(cats) else ('', '', '', '')
        put(ws, f'{C_SEQ}{r}', f'=IF({C_NAME}{r}="","",ROW()-{CAT_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{C_NAME}{r}', c[0] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{C_DIR}{r}', c[1] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{C_TO}{r}', c[2] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{C_NOTE}{r}', c[3] or None, F_NOTE, FILL_IN, align=ALW)
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 15)
    add_list_dv(ws, f'{C_DIR}{CAT_R0}:{C_DIR}{CAT_R1}', '"收,支,收支"')
    add_list_dv(ws, f'{C_TO}{CAT_R0}:{C_TO}{CAT_R1}', '"' + ','.join(TO_ALL) + '"')

    # ── 费用项目 AH:AK ──
    _block_title(ws, E_SEQ, E_NOTE, '费用项目（资金台帐「费用项目」下拉）', 'FFC65911')
    header(ws, 4, [(E_SEQ, '序号'), (E_NAME, '费用项目'), (E_CLASS, '利润表归类'), (E_NOTE, '说明')],
           'FFC65911', height=34)
    exps = ctx['expense_items']
    for i in range(EXP_R1 - EXP_R0 + 1):
        r = EXP_R0 + i
        e = exps[i] if i < len(exps) else ('', '', '')
        put(ws, f'{E_SEQ}{r}', f'=IF({E_NAME}{r}="","",ROW()-{EXP_R0 - 1})', F_AUTO, FILL_AUTO, align=AC)
        put(ws, f'{E_NAME}{r}', e[0] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{E_CLASS}{r}', e[1] or None, F_IN, FILL_IN, align=AC)
        put(ws, f'{E_NOTE}{r}', e[2] or None, F_NOTE, FILL_IN, align=ALW)
    add_list_dv(ws, f'{E_CLASS}{EXP_R0}:{E_CLASS}{EXP_R1}', '"' + ','.join(EC_ALL) + '"')

    # ── 报损月表清单 AM:AO ──
    _block_title(ws, M_MONTH, M_NOTE, '报损月表清单（几月的报损记在哪张表）', 'FF833C0C')
    header(ws, 4, [(M_MONTH, '月份'), (M_SHEET, '报损表名'), (M_NOTE, '说明')], 'FF833C0C', height=34)
    bsm = ctx['baosun_months']      # {月份: 表名}
    for m in range(1, 13):
        r = BSM_R0 + m - 1
        put(ws, f'{M_MONTH}{r}', m, F_TXTB, FILL_AUTO, '0"月"', AC)
        put(ws, f'{M_SHEET}{r}', bsm.get(m), F_IN, FILL_IN, align=AC)
        put(ws, f'{M_NOTE}{r}',
            f'=IF({M_SHEET}{r}="","（这个月没有报损表）",IF(ISREF(INDIRECT("\'"&{M_SHEET}{r}&"\'!A1")),"√ 找到这张表","✗ 没有叫这个名字的表"))',
            F_NOTE, FILL_AUTO, align=AL)
    put(ws, f'{M_MONTH}{BSM_R1 + 2}',
        '新的一个月：右键上个月的报损表→移动或复制→建立副本，改成「2026年10月报损」这样的表名，清掉数量，再回这里填上表名。',
        F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'{M_MONTH}{BSM_R1 + 2}:{M_NOTE}{BSM_R1 + 4}')

    widths(ws, {'A': 22, 'B': 30, 'C': 2, G_SEQ: 5, G_NAME: 16, G_SUP: 13, G_BSNAME: 18, G_SPEC: 7,
                G_PACK: 6, G_COST: 10, G_PRICE: 10, G_Q0: 9, G_P0: 9, G_A0: 11, 'O': 2,
                S_SEQ: 5, S_NAME: 16, S_AP0: 15, S_CONTACT: 15, S_NOTE: 16, 'U': 2,
                A_SEQ: 5, A_NAME: 14, A_BAL0: 14, A_BALNOW: 15, A_NOTE: 16, 'AA': 2,
                C_SEQ: 5, C_NAME: 12, C_DIR: 6, C_TO: 9, C_NOTE: 40, 'AG': 2,
                E_SEQ: 5, E_NAME: 13, E_CLASS: 11, E_NOTE: 28, 'AL': 2,
                M_MONTH: 7, M_SHEET: 17, M_NOTE: 22})
    ws.freeze_panes = 'A5'

    # ── 定义名称 ──
    def dn(name, ref):
        wb.defined_names[name] = DefinedName(name, attr_text=ref)
    dn('年度', f'{SH_BASE}!$B$5')
    dn('年初日', f'{SH_BASE}!$B$7')
    dn('年末日', f'{SH_BASE}!$B$8')
    # 下拉清单的长度 = 到最后一个非空格为止（中间空了一格也不会把后面的截掉）
    def lst(sh, col, r0, r1):
        rng = f'{sh}!${col}${r0}:${col}${r1}'
        return f'OFFSET({sh}!${col}${r0},0,0,MAX(1,IFERROR(LOOKUP(2,1/({rng}<>""),ROW({rng}))-{r0 - 1},1)),1)'
    dn('客户列表', lst(SH_OV, 'B', OV_R0, OV_R1))
    dn('商品列表', lst(SH_BASE, G_NAME, BASE_R0, BASE_R1))
    dn('供应商列表', lst(SH_BASE, S_NAME, SUP_R0, SUP_R1))
    dn('账户列表', lst(SH_BASE, A_NAME, ACC_R0, ACC_R1))
    dn('收支项目列表', lst(SH_BASE, C_NAME, CAT_R0, CAT_R1))
    dn('费用项目列表', lst(SH_BASE, E_NAME, EXP_R0, EXP_R1))
    return ws
