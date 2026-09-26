# -*- coding: utf-8 -*-
"""A064 饮料经销财务模板 · 共用常量（表名 / 行列地址 / 容量）与样式。

这个文件就是各张表之间的「接口约定」：
凡是一张表要引用另一张表的某一列、某一段行，一律从这里取常量，
不在各模块里手写地址，免得一处改了另一处没跟上。
"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

# ─────────────────────────── 表名 ───────────────────────────
SH_HOME = '首页'
SH_BASE = '基础资料'
SH_CASH = '资金台帐'
SH_BUY  = '采购进货'
SH_FEE  = '全量费用明细总表'  # 你补的费用报销明细：现在不用录，从资金台帐的费用支出自动列出（原来的列不变）
SH_QR   = '收款码到账对账'    # 从资金台帐自动生成：收款码每天收款 vs 转到银行的到账，差额是手续费
SH_OV   = '总览汇总'      # 原表
SH_OUT  = '出库明细'      # 原表
SH_RP   = '收付款明细'    # 原表（改成从资金台帐自动提取）
SH_CUN  = '存条明细'      # 原表
SH_INV  = '商品库存'      # 原表（行改成自动列出）
SH_Q    = '客户查询'      # 原表（修好）
SH_RPS  = '收付款汇总'
SH_CST  = '客户对账单'
SH_SST  = '供应商对账单'
SH_STK  = '公司库存'
SH_COST = '成本计算'
SH_BSS  = '报损汇总一览'  # 原表2
SH_BSL  = '报损明细台账'  # 原表2：原来一个月一张，合成一年一张（加「月份」列，按月筛选）
SH_EXP  = '费用汇总'
SH_PL   = '利润表'
SH_BAL  = '资产负债表'
SH_CHK  = '数据校验'
SH_AUX  = '_辅助'

SHEET_ORDER = [SH_HOME, SH_BASE, SH_CASH, SH_BUY, SH_OV, SH_OUT, SH_RP, SH_CUN, SH_INV, SH_Q,
               SH_FEE, SH_QR, SH_RPS, SH_CST, SH_SST, SH_STK, SH_COST, SH_BSL, SH_BSS,
               SH_EXP, SH_PL, SH_BAL, SH_CHK, SH_AUX]

def q(sh):
    """跨表引用时的表名写法：带数字开头或特殊字符的表名要加单引号"""
    return f"'{sh}'" if (sh[0].isdigit() or any(ch in sh for ch in ' -()（）')) else sh

# ─────────────────────────── 基础资料 ───────────────────────────
# 参数块 A:B
P_YEAR   = f'{SH_BASE}!$B$5'     # 会计年度
P_NAME   = f'{SH_BASE}!$B$6'     # 公司名称
P_START  = f'{SH_BASE}!$B$7'     # 年初日 =DATE(年度,1,1)
P_END    = f'{SH_BASE}!$B$8'     # 年末日
P_LOAN0  = f'{SH_BASE}!$B$9'     # 年初借款余额（手填，默认 0）
P_OTH0   = f'{SH_BASE}!$B$10'    # 年初其他应收(+)/应付(-) 净额（手填，默认 0）

BASE_R0, BASE_R1 = 5, 104        # 商品档案 100 行
# 商品档案 D:N
G_SEQ, G_NAME, G_SUP, G_BSNAME, G_SPEC, G_PACK, G_COST, G_PRICE, G_Q0, G_P0, G_A0 = \
    'D', 'E', 'F', 'G', 'H', 'I', 'J', 'K', 'L', 'M', 'N'
SUP_R0, SUP_R1 = 5, 54           # 供应商档案 50 行  P:T
S_SEQ, S_NAME, S_AP0, S_CONTACT, S_NOTE = 'P', 'Q', 'R', 'S', 'T'
ACC_R0, ACC_R1 = 5, 24           # 资金账户 20 行  V:Z
A_SEQ, A_NAME, A_BAL0, A_BALNOW, A_NOTE = 'V', 'W', 'X', 'Y', 'Z'
CAT_R0, CAT_R1 = 5, 34           # 收支项目 30 行  AB:AF
C_SEQ, C_NAME, C_DIR, C_TO, C_NOTE = 'AB', 'AC', 'AD', 'AE', 'AF'
EXP_R0, EXP_R1 = 5, 54           # 费用项目 50 行  AH:AK
E_SEQ, E_NAME, E_CLASS, E_NOTE = 'AH', 'AI', 'AJ', 'AK'

def base_rng(col, r0, r1):
    return f"{SH_BASE}!${col}${r0}:${col}${r1}"

GOODS_NAMES = base_rng(G_NAME, BASE_R0, BASE_R1)
SUP_NAMES   = base_rng(S_NAME, SUP_R0, SUP_R1)
ACC_NAMES   = base_rng(A_NAME, ACC_R0, ACC_R1)
CAT_NAMES   = base_rng(C_NAME, CAT_R0, CAT_R1)
EXP_NAMES   = base_rng(E_NAME, EXP_R0, EXP_R1)

# 「去向」代码：收支项目决定这笔钱在报表里落到哪
TO_AR    = '应收'       # 客户回款 / 客户退款：冲减客户未付货款
TO_AP    = '应付'       # 供应商付款 / 供应商退款：冲减应付供应商
TO_EXP   = '费用'       # 费用支出：按费用项目进利润表
TO_REB   = '返利收入'   # 厂家返利 / 补贴：利润表·其他业务收入
TO_OI    = '其他收入'   # 营业外收入
TO_INV   = '投入'       # 股东投入：所有者权益
TO_DRAW  = '提取'       # 股东提取：所有者权益
TO_LOAN  = '借款'       # 借入 / 归还：短期借款
TO_XFER  = '转账'       # 账户之间互转：不影响报表，收支要成对
TO_OTH   = '其他往来'   # 押金等：其他应收 / 其他应付
TO_ALL = [TO_AR, TO_AP, TO_EXP, TO_REB, TO_OI, TO_INV, TO_DRAW, TO_LOAN, TO_XFER, TO_OTH]

# 费用「利润表归类」
EC_TAX, EC_SELL, EC_ADMIN, EC_FIN, EC_NOI, EC_IT = '税金及附加', '销售费用', '管理费用', '财务费用', '营业外支出', '所得税费用'
EC_ALL = [EC_SELL, EC_ADMIN, EC_FIN, EC_TAX, EC_NOI, EC_IT]

# ─────────────────────────── 资金台帐 ───────────────────────────
CASH_HDR = 6
CASH_R0, CASH_R1 = 7, 3006       # 3000 行
(K_SEQ, K_DATE, K_ACC, K_CAT, K_MEMO, K_IN, K_OUT, K_ABAL, K_TBAL,
 K_CUS, K_SUP, K_EXP, K_WHO, K_NOTE, K_CHK) = 'A B C D E F G H I J K L M N O'.split()
K_LAST = K_CHK                   # 看得见的最后一列
# 隐藏辅助列
K_CAT2, K_TO, K_CSEQ, K_SSEQ, K_NET, K_ESEQ = 'P', 'Q', 'R', 'S', 'T', 'U'
# P 实际收支项目（没选项目时按客户/供应商/费用推断） Q 去向  R 客户收款排序键  S 供应商付款排序键  T 净额=收入-支出
# U 费用排序键（给【全量费用明细总表】按日期列出费用用）

def cash(col, absolute=True):
    return f"{SH_CASH}!${col}${CASH_R0}:${col}${CASH_R1}" if absolute else f"{SH_CASH}!{col}{CASH_R0}:{col}{CASH_R1}"

# ─────────────────────────── 采购进货 ───────────────────────────
BUY_HDR = 3
BUY_R0, BUY_R1 = 4, 1503         # 1500 行
(B_SEQ, B_DATE, B_SUP, B_GOODS, B_QTY, B_PRICE, B_AMT, B_GIFT, B_NOTE, B_CHK) = 'A B C D E F G H I J'.split()

def buy(col):
    return f"{SH_BUY}!${col}${BUY_R0}:${col}${BUY_R1}"

# ─────────────────────────── 全量费用明细总表（自动从资金台帐列出，列跟你原来的表一样） ───────────────────────────
FEE_HDR = 4
FEE_R0, FEE_R1 = 5, 2004         # 2000 行
(F_SEQ, F_DATE, F_CAT, F_DETAIL, F_WHO, F_AMT, F_ACC, F_REMK, F_POS) = 'A B C D E F G H I'.split()   # I 隐藏：资金台帐第几行

# ─────────────────────────── 报损明细台账（一年一张） ───────────────────────────
BSL_HDR = 5
BSL_R0, BSL_R1 = 6, 2005         # 2000 行
(L_SEQ, L_MON, L_COMP, L_NAME, L_SPEC, L_QTY, L_UNIT, L_PRICE, L_AMT, L_WHY, L_NOTE, L_CHK) = \
    'A B C D E F G H I J K L'.split()

def bsl(col):
    return f"{SH_BSL}!${col}${BSL_R0}:${col}${BSL_R1}"

# ─────────────────────────── 出库明细（原表） ───────────────────────────
OUT_HDR = 3
OUT_R0, OUT_R1 = 4, 3003         # 3000 行（原来到 545）
(O_SEQ, O_DATE, O_CUS, O_GOODS, O_QTY, O_PRICE, O_AMT, O_NOTE) = 'A B C D E F G H'.split()

def out(col):
    return f"{SH_OUT}!${col}${OUT_R0}:${col}${OUT_R1}"

# ─────────────────────────── 存条明细（原表） ───────────────────────────
CUN_HDR = 3
CUN_R0, CUN_R1 = 4, 1003         # 1000 行
(U_SEQ, U_DATE, U_CUS, U_GOODS, U_QTY, U_PRICE, U_AMT, U_NOTE) = 'A B C D E F G H'.split()

def cun(col):
    return f"{SH_CUN}!${col}${CUN_R0}:${col}${CUN_R1}"

# ─────────────────────────── 总览汇总（原表，客户名单就在这里） ───────────────────────────
OV_TOP = 4                       # 第 4 行是合计（引用底部合计行）
OV_R0, OV_R1 = 5, 304            # 300 个客户
OV_TOT = 305
(V_SEQ, V_NAME, V_DEP, V_USED, V_LEFT, V_PAID, V_OWE, V_STAT, V_IQ, V_IA, V_OPEN) = 'A B C D E F G H I J K'.split()
CUS_NAMES = f"{SH_OV}!$B${OV_R0}:$B${OV_R1}"
CUS_OPEN  = f"{SH_OV}!$K${OV_R0}:$K${OV_R1}"

# ─────────────────────────── 商品库存（原表，行改为自动列出） ───────────────────────────
INV_HDR = 3
INV_R0, INV_R1 = 4, 603          # 600 个「客户+商品」组合
INV_TOT = 604

# ─────────────────────────── 收付款明细（原表，改为自动提取） ───────────────────────────
RP_HDR = 3
RP_R0, RP_R1 = 4, 1003           # 客户收款 1000 行（A:F，G 隐藏存行号）
RP_SR1 = 503                     # 供应商付款 500 行（I:N，O 隐藏存行号）

# ─────────────────────────── 成本计算 ───────────────────────────
COST_R0, COST_R1 = 8, 107        # 与商品档案一一对应（基础资料 5..104 → 8..107）
COST_TOT = 108
COST_FIX = 7                     # A..G 固定列：序号 商品 报损品名 每件数量 参考进价 期初数量 期初金额
COST_FIELDS = ['采购数量', '采购金额', '可用数量', '可用金额', '平均单价', '出库数量', '出库金额',
               '出库成本', '报损数量', '报损金额', '结存数量', '结存金额', '估算成本']
COST_W = len(COST_FIELDS)        # 每个月 13 列

def cost_col(m, field):
    """第 m 月（1..12）某字段所在列字母"""
    return CL(COST_FIX + (m - 1) * COST_W + COST_FIELDS.index(field) + 1)

# ─────────────────────────── 月份工具 ───────────────────────────
YEAR = '年度'                    # 定义名称 → 基础资料!$B$5

def d_from(m):                   # 该月第一天（m 可以是数字或单元格引用）
    return f'DATE({YEAR},{m},1)'

def d_to_excl(m):                # 下月第一天（不含）
    return f'DATE({YEAR},{m}+1,1)'

def month_crit(date_rng, m):
    """SUMIFS 用的「日期在第 m 月」两对条件"""
    return f'{date_rng},">="&{d_from(m)},{date_rng},"<"&{d_to_excl(m)}'

# ─────────────────────────── 样式（照原表：微软雅黑） ───────────────────────────
YH = '微软雅黑'
F_TITLE = Font(name=YH, sz=16, bold=True, color='FFFFFFFF')
F_TIP   = Font(name=YH, sz=10, color='FFFF6600')
F_HDR   = Font(name=YH, sz=11, bold=True, color='FFFFFFFF')
F_HDR_D = Font(name=YH, sz=11, bold=True, color='FF1F3864')
F_TXT   = Font(name=YH, sz=10, color='FF000000')
F_TXTB  = Font(name=YH, sz=10, bold=True, color='FF000000')
F_IN    = Font(name=YH, sz=10, color='FF1F4E79')
F_AUTO  = Font(name=YH, sz=10, color='FF404040')
F_NOTE  = Font(name=YH, sz=9, color='FF808080')
F_RED   = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_KPI_L = Font(name=YH, sz=11, bold=True, color='FF1F3864')
F_KPI_V = Font(name=YH, sz=13, bold=True, color='FFC00000')
F_SEC   = Font(name=YH, sz=12, bold=True, color='FF000000')
F_BIG   = Font(name=YH, sz=12, bold=True, color='FF000000')

def fill(rgb):
    return PatternFill('solid', fgColor=rgb)

# 原表各张的标题色（标签颜色也用同一个色，一眼能认出是哪张）
C_OV, C_OUT, C_RP, C_CUN, C_INV, C_Q = 'FFC00000', 'FF4472C4', 'FF70AD47', 'FF0070C0', 'FF5B9BD5', 'FFBF8F00'
C_BS, C_FEE = 'FF833C0C', 'FFC65911'
# 新表标题色
C_CASH, C_BUY, C_BASE, C_RPT, C_STMT, C_CHK, C_COST = 'FF2F75B5', 'FF7030A0', 'FF595959', 'FF375623', 'FFBF8F00', 'FF833C0C', 'FF808080'
C_INCOME, C_EXPENSE = 'FF00B050', 'FFFFC000'      # 资金台帐：收入绿、支出黄（照你给的截图）

FILL_TIP  = fill('FFFFF2CC')
FILL_IN   = fill('FFFFF7E0')   # 手工录入：淡黄
FILL_AUTO = fill('FFF2F2F2')   # 自动计算：淡灰
FILL_KPI  = fill('FFD9E1F2')   # 资金台帐顶部汇总条（截图里的淡蓝灰）
FILL_BAND = fill('FF8EA9DB')   # 截图第 4 行那条蓝带
FILL_SUBH = fill('FFDDEBF7')   # 次级表头 / 小计
FILL_TOT  = fill('FFFCE4D6')   # 合计行（原总览汇总的淡橙）
FILL_OK   = fill('FFE2EFDA')
FILL_WARN = fill('FFFFC7CE')
FILL_SEL  = fill('FFFFFF00')   # 查询/对账单的选择格（原客户查询的亮黄）

thin = Side(style='thin', color='FFBFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBD = Border()

AC  = Alignment(horizontal='center', vertical='center')
ACW = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL  = Alignment(horizontal='left', vertical='center')
ALW = Alignment(horizontal='left', vertical='center', wrap_text=True)
AR_ = Alignment(horizontal='right', vertical='center')

MONEY = '\\¥#,##0.00'                      # 原表用的格式
MONEY2 = '#,##0.00;[Red]-#,##0.00;"-"'     # 报表用：负数红、零显示 -
QTY   = '[=0]"-";General'                  # 数量：整数不带小数点，半件照样显示 17.5，零显示 -
PRICE = '#,##0.00##'
DATE  = 'yyyy/mm/dd'
PCT   = '0.0%;[Red]-0.0%;"-"'


def at(rng, k):
    """按位置取第 k 格：INDEX(区域,k)。辅助公式一律按位置取，用户在明细表里插行/删行也不会串位或变 #REF!"""
    return f'INDEX({rng},{k})'


def upto(rng, k):
    """区域开头到第 k 格：$C$4:INDEX($C$4:$C$1003,k)（累计 COUNTIFS 用）"""
    first = rng.split(':')[0]
    return f'{first}:INDEX({rng},{k})'


def put(ws, coord, value=None, font=None, fill_=None, fmt=None, align=None, border=True):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if fill_ is not None:
        c.fill = fill_
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border:
        c.border = BD
    return c


def title(ws, text, last_col, color, tip=None, h1=33, h2=None):
    """第 1 行标题带 + 第 2 行 💡 提示（和原表一个样）；提示长就把第 2 行加高，免得被截掉"""
    if h2 is None:
        width = sum((ws.column_dimensions[CL(i)].width or 9) for i in range(1, CI(last_col) + 1))
        per_line = max(20, int(width / 1.9))          # 一行大约能放多少个汉字（10 号雅黑）
        lines = -(-len(tip or '') // per_line)
        h2 = max(24, 16 * lines + 6)
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', text, F_TITLE, fill(color), align=AC, border=False)
    ws.row_dimensions[1].height = h1
    if tip:
        ws.merge_cells(f'A2:{last_col}2')
        put(ws, 'A2', tip, F_TIP, FILL_TIP, align=ALW, border=False)
        ws.row_dimensions[2].height = h2


def header(ws, row, cols_texts, color, font=F_HDR, height=37):
    """cols_texts: [(列字母, 文字), ...]"""
    for col, t in cols_texts:
        put(ws, f'{col}{row}', t, font, fill(color), align=ACW)
    ws.row_dimensions[row].height = height


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def link(cell, sheet, ref='A1'):
    """表内跳转链接：写成 location（Excel 自己存内部链接的方式），WPS / Excel / 手机版都能点过去。
    不能写成 '#表名!A1' 的外部链接——WPS 不认，点了没反应。"""
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{ref}",
                               display=None if cell.value is None else str(cell.value))
    return cell


def style_rows(ws, r0, r1, cols, auto=(), fmts=None, aligns=None, bold=(), height=None):
    """数据区统一样式（跟 A064 新表一样）：白底细灰框；手填列黑字，公式列灰底灰字；不再隔行涂色。
    cols: 列字母串；auto: 公式列；fmts / aligns: {列: 格式 / 对齐}；bold: 要加粗的列"""
    fmts, aligns = fmts or {}, aligns or {}
    for r in range(r0, r1 + 1):
        for col in cols:
            c = ws[f'{col}{r}']
            is_auto = col in auto
            c.font = (F_TXTB if col in bold else (F_AUTO if is_auto else F_TXT))
            c.fill = FILL_AUTO if is_auto else PatternFill(fill_type=None)
            c.border = BD
            c.alignment = aligns.get(col, AC)
            if col in fmts:
                c.number_format = fmts[col]
        if height:
            ws.row_dimensions[r].height = height


def add_date_dv(ws, sqref):
    """日期列：不在本会计年度就弹提醒（只提醒不拦，跨年补录也能记）"""
    dv = DataValidation(type='date', operator='between', formula1='年初日', formula2='年末日', allow_blank=True,
                        showErrorMessage=True, errorStyle='warning')
    dv.errorTitle, dv.error = '日期不对', '日期不在本会计年度，或者不是真日期（要像 2026/9/25 这样录）'
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def add_list_dv(ws, sqref, formula, prompt=None, allow_blank=True, stop=True):
    dv = DataValidation(type='list', formula1=formula, allow_blank=allow_blank,
                        showErrorMessage=True, errorStyle='stop' if stop else 'warning')
    if prompt:
        dv.promptTitle, dv.prompt = '提示', prompt
        dv.showInputMessage = True
    dv.error = '请从下拉里选；清单里没有的，先到【基础资料】/【总览汇总】里登记'
    dv.errorTitle = '不在清单里'
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv
