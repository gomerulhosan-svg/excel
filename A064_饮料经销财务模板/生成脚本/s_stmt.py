# -*- coding: utf-8 -*-
"""【收付款汇总】【客户对账单】【供应商对账单】（版式参考原【客户查询】：黄格子选人，💰/📦/📋 分块）。

对账单的明细要把两张表（出库+资金台帐 / 采购+资金台帐）的行按日期混排，
不用 FILTER/SORT：在 _辅助 里给每一行算一个「日期＋行号/10万」的排序键（不符合条件的留空），
对账单第 k 行用 SMALL 取第 k 小的键，再 MATCH 回原行。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font

from common import *

N_OUT = OUT_R1 - OUT_R0 + 1        # 3000
N_BUY = BUY_R1 - BUY_R0 + 1        # 1500
N_CASH = CASH_R1 - CASH_R0 + 1     # 3000
AUX_C0, AUX_C1 = 4, 4 + N_OUT + N_CASH - 1       # F 列：客户对账单排序键 4..6003
AUX_S0, AUX_S1 = 4, 4 + N_BUY + N_CASH - 1       # H 列：供应商对账单排序键 4..4503
AUX_CST_INV = 'E'                                # E 列：客户对账单·存条商品取数

ST_DET0 = 39                       # 对账单明细第一行
ST_N = 300                         # 最多 300 行
ST_DET1 = ST_DET0 + ST_N - 1
ST_TOT = ST_DET1 + 1


def _sel_bar(ws, who_label, dv_name, ctx):
    ws.row_dimensions[4].height = 29
    put(ws, 'A4', who_label, Font(name=YH, sz=12, bold=True), align=AR_, border=False)
    ws.merge_cells('B4:D4')
    put(ws, 'B4', ctx.get('default_' + dv_name), Font(name=YH, sz=12, bold=True), FILL_SEL, align=AC)
    add_list_dv(ws, 'B4', f'={dv_name}', '下拉选')
    put(ws, 'F4', '起始日期：', Font(name=YH, sz=11, bold=True), align=AR_, border=False)
    put(ws, 'G4', '=年初日', Font(name=YH, sz=11, bold=True), FILL_SEL, DATE, AC)
    put(ws, 'H4', '截止日期：', Font(name=YH, sz=11, bold=True), align=AR_, border=False)
    put(ws, 'I4', '=年末日', Font(name=YH, sz=11, bold=True), FILL_SEL, DATE, AC)
    put(ws, 'A5', '对账单位：', F_TXTB, align=AR_, border=False)
    put(ws, 'B5', f'=IF({P_NAME}="","（在基础资料填公司名称）",{P_NAME})', F_TXT, align=AL, border=False)
    ws.merge_cells('B5:D5')
    put(ws, 'F5', '打印日期：', F_TXTB, align=AR_, border=False)
    put(ws, 'G5', '=TODAY()', F_TXT, fmt=DATE, align=AC, border=False)


def _section(ws, row, text, color, c1='J'):
    ws.merge_cells(f'A{row}:{c1}{row}')
    put(ws, f'A{row}', text, Font(name=YH, sz=12, bold=True, color='FFFFFFFF'), fill(color), align=AL)
    ws.row_dimensions[row].height = 26


# ───────────────────────── _辅助：排序键 ─────────────────────────
def build_aux_stmt(wb, ctx):
    ws = wb[SH_AUX]
    put(ws, f'{AUX_CST_INV}3', '客户对账单：存条第几个', F_NOTE, border=False)
    put(ws, 'F3', '客户对账单排序键', F_NOTE, border=False)
    put(ws, 'H3', '供应商对账单排序键', F_NOTE, border=False)
    C, S = SH_CST, SH_SST
    rng = lambda sh, d: f'ISNUMBER({d}),{d}>={sh}!$G$4,{d}<{sh}!$I$4+1'
    # 全部按位置取（INDEX(区域,k)）：明细表插行/删行不会串位
    IB = f'{SH_INV}!$B${INV_R0}:$B${INV_R1}'
    for k in range(1, INV_R1 - INV_R0 + 2):
        r = 3 + k
        b = at(IB, k)
        ws[f'{AUX_CST_INV}{r}'] = f'=IF(AND({b}<>"",{b}={C}!$B$4,N({C}!$K$12)>0),COUNTIF({upto(IB, k)},{C}!$B$4),"")'
    for k in range(1, N_OUT + 1):
        r = AUX_C0 + k - 1
        cus, d = at(out(O_CUS), k), at(out(O_DATE), k)
        ws[f'F{r}'] = f'=IF(AND({cus}<>"",{cus}={C}!$B$4,{rng(C, d)}),{d}+ROW()/100000,"")'
    for k in range(1, N_CASH + 1):
        r = AUX_C0 + N_OUT + k - 1
        cus, to, d = at(cash(K_CUS), k), at(cash(K_TO), k), at(cash(K_DATE), k)
        ws[f'F{r}'] = f'=IF(AND({cus}<>"",{cus}={C}!$B$4,{to}="{TO_AR}",{rng(C, d)}),{d}+ROW()/100000,"")'
    for k in range(1, N_BUY + 1):
        r = AUX_S0 + k - 1
        sup, d = at(buy(B_SUP), k), at(buy(B_DATE), k)
        ws[f'H{r}'] = f'=IF(AND({sup}<>"",{sup}={S}!$B$4,{rng(S, d)}),{d}+ROW()/100000,"")'
    for k in range(1, N_CASH + 1):
        r = AUX_S0 + N_BUY + k - 1
        sup, to, d = at(cash(K_SUP), k), at(cash(K_TO), k), at(cash(K_DATE), k)
        ws[f'H{r}'] = f'=IF(AND({sup}<>"",{sup}={S}!$B$4,{to}="{TO_AP}",{rng(S, d)}),{d}+ROW()/100000,"")'
    put(ws, 'F2', f'=COUNT(F{AUX_C0}:F{AUX_C1})', F_NOTE, border=False)       # 本张客户对账单的往来笔数
    put(ws, 'H2', f'=COUNT(H{AUX_S0}:H{AUX_S1})', F_NOTE, border=False)


# ───────────────────────── 客户对账单 ─────────────────────────
def build_cst(wb, ctx):
    ws = wb.create_sheet(SH_CST)
    ws.sheet_properties.tabColor = C_STMT[2:]
    title(ws, '客户对账单', 'J', C_STMT,
          '💡 在黄色格子选客户、起止日期，自动出对账单：货款汇总、存条商品结余、逐笔往来（出库领用和收款按日期排好）。'
          '直接打印给客户签字即可。数据来自【出库明细】【资金台帐】【商品库存】。')
    _sel_bar(ws, '🔍 客户名称：', '客户列表', ctx)
    sel, d0, d1 = '$B$4', '$G$4', '$I$4'
    cus = f'{SH_OV}!$B${OV_R0}:$B${OV_R1}'
    opening = (f'SUMIF({cus},{sel},{CUS_OPEN})+SUMIFS({out(O_AMT)},{out(O_CUS)},{sel},{out(O_DATE)},">="&年初日,{out(O_DATE)},"<"&{d0})'
               f'-SUMIFS({cash(K_NET)},{cash(K_CUS)},{sel},{cash(K_TO)},"{TO_AR}",{cash(K_DATE)},">="&年初日,{cash(K_DATE)},"<"&{d0})')
    # 💰 货款汇总
    _section(ws, 7, '💰 货款汇总', 'FF70AD47')
    heads = ['期初欠款', '本期领用金额', '本期收款', '期末欠款', '状态', '存条剩余数量(件)', '存条剩余金额']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}8', t, F_HDR, fill('FF70AD47'), align=ACW)
    ws.row_dimensions[8].height = 30
    vals = [f'={opening}',
            f'=SUMIFS({out(O_AMT)},{out(O_CUS)},{sel},{out(O_DATE)},">="&{d0},{out(O_DATE)},"<"&({d1}+1))',
            f'=SUMIFS({cash(K_NET)},{cash(K_CUS)},{sel},{cash(K_TO)},"{TO_AR}",{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&({d1}+1))',
            '=A9+B9-C9',
            '=IF(ROUND(D9,2)=0,"已结清",IF(D9>0,"有欠款","预存款"))',
            f'=IF(N($K$12)>0,SUMIF({SH_INV}!$B${INV_R0}:$B${INV_R1},{sel},{SH_INV}!$I${INV_R0}:$I${INV_R1}),"无存条")',
            f'=IF(N($K$12)>0,SUMIF({SH_INV}!$B${INV_R0}:$B${INV_R1},{sel},{SH_INV}!$J${INV_R0}:$J${INV_R1}),"")']
    for i, v in enumerate(vals):
        put(ws, f'{CL(i + 1)}9', v, Font(name=YH, sz=11, bold=True), FILL_OK,
            QTY if i == 5 else (None if i == 4 else MONEY2), AC)
    ws.row_dimensions[9].height = 24
    put(ws, 'A10', '说明：欠款为正＝客户还欠我们；为负＝客户预存（存条款没用完）。', F_NOTE, align=AL, border=False)
    ws.merge_cells('A10:J10')
    put(ws, 'A11', '以上往来核对无误。客户签字（盖章）：', F_TXTB, align=AL, border=False)
    ws.merge_cells('A11:E11')
    put(ws, 'G11', '日期：', F_TXTB, align=AR_, border=False)
    ws.row_dimensions[11].height = 30

    # 📦 存条商品结余
    _section(ws, 13, '📦 存条商品结余（截至目前，来自【商品库存】）', 'FF5B9BD5')
    heads = ['商品品类', '累计存条数量', '累计存条金额', '已领用数量', '已领用金额', '剩余数量', '剩余金额']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}14', t, F_HDR, fill('FF5B9BD5'), align=ACW)
    src_cols = ['C', 'D', 'F', 'G', 'H', 'I', 'J']
    aux_rng = f'{SH_AUX}!${AUX_CST_INV}$4:${AUX_CST_INV}${4 + INV_R1 - INV_R0}'
    for k in range(1, 21):
        r = 14 + k
        put(ws, f'K{r}', f'=IFERROR(MATCH({k},{aux_rng},0),"")', F_NOTE, border=False)
        for j, sc in enumerate(src_cols):
            put(ws, f'{CL(j + 1)}{r}', f'=IF($K{r}="","",INDEX({SH_INV}!${sc}${INV_R0}:${sc}${INV_R1},$K{r}))', F_TXT,
                fmt=None if j == 0 else (MONEY2 if j in (2, 4, 6) else QTY), align=AC)
    put(ws, 'K12', f'=SUMIF({SH_INV}!$B${INV_R0}:$B${INV_R1},{sel},{SH_INV}!$D${INV_R0}:$D${INV_R1})', F_NOTE, border=False)
    put(ws, 'A35', (f'=IF(N($K$12)<=0,"这个客户没有存条，都是直接领货，这一块不适用",'
                    f'IF(COUNT({aux_rng})>20,"⚠ 这个客户的存条商品超过 20 种，只列了前 20 种，完整的看【客户查询】",""))'),
        F_RED, align=AL, border=False)
    ws.merge_cells('A35:J35')

    # 📋 往来明细
    _section(ws, 37, '📋 往来明细（按日期排）', 'FF4472C4')
    heads = ['序号', '日期', '摘要', '商品品类', '数量(件)', '单价', '领用金额', '收款金额', '欠款余额', '备注']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}38', t, F_HDR, fill('FF4472C4'), align=ACW)
    ws.row_dimensions[38].height = 28
    keys = f'{SH_AUX}!$F${AUX_C0}:$F${AUX_C1}'
    o = lambda col: f'{SH_OUT}!${col}${OUT_R0}:${col}${OUT_R1}'
    c = lambda col: f'{SH_CASH}!${col}${CASH_R0}:${col}${CASH_R1}'
    # 期初行
    r = ST_DET0
    put(ws, f'B{r}', f'={d0}', F_TXT, fmt=DATE, align=AC)
    put(ws, f'C{r}', '期初欠款', F_TXTB, align=AC)
    put(ws, f'I{r}', '=A9', F_TXTB, fmt=MONEY2, align=AR_)
    for col in 'ADEFGHJ':
        put(ws, f'{col}{r}', None)
    for k in range(1, ST_N):
        r = ST_DET0 + k
        K = f'$K{r}'
        isout = f'{K}<={N_OUT}'
        ci = f'{K}-{N_OUT}'
        put(ws, f'K{r}', f'=IF({k}>{SH_AUX}!$F$2,"",IFERROR(MATCH(SMALL({keys},{k}),{keys},0),""))', F_NOTE, border=False)
        put(ws, f'A{r}', f'=IF({K}="","",{k})', F_TXT, align=AC)
        put(ws, f'B{r}', f'=IF({K}="","",IF({isout},INDEX({o("B")},{K}),INDEX({c(K_DATE)},{ci})))', F_TXT, fmt=DATE, align=AC)
        put(ws, f'C{r}', (f'=IF({K}="","",IF({isout},"出库领用",IF(INDEX({c(K_NET)},{ci})>=0,"收款","退款")'
                          f'&"（"&INDEX({c(K_ACC)},{ci})&"）"))'), F_TXT, align=AC)
        put(ws, f'D{r}', f'=IF({K}="","",IF({isout},INDEX({o("D")},{K})&"",""))', F_TXT, align=AC)
        put(ws, f'E{r}', f'=IF({K}="","",IF({isout},INDEX({o("E")},{K}),""))', F_TXT, fmt=QTY, align=AC)
        put(ws, f'F{r}', f'=IF({K}="","",IF({isout},IF(INDEX({o("F")},{K})="","",INDEX({o("F")},{K})),""))', F_TXT, fmt=PRICE, align=AC)
        put(ws, f'G{r}', f'=IF({K}="","",IF({isout},N(INDEX({o("G")},{K})),""))', F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'H{r}', f'=IF({K}="","",IF({isout},"",INDEX({c(K_NET)},{ci})))', F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'I{r}', f'=IF({K}="","",I{r - 1}+N(G{r})-N(H{r}))', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'J{r}', f'=IF({K}="","",IF({isout},INDEX({o("H")},{K})&"",INDEX({c(K_MEMO)},{ci})&""))', F_NOTE, align=AL)
    r = ST_TOT
    put(ws, f'C{r}', '本期合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'E{r}', f'=SUM(E{ST_DET0}:E{ST_DET1})', F_TXTB, FILL_TOT, QTY, AC)
    put(ws, f'G{r}', f'=SUM(G{ST_DET0}:G{ST_DET1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'H{r}', f'=SUM(H{ST_DET0}:H{ST_DET1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'I{r}', '=D9', F_TXTB, FILL_TOT, MONEY2, AR_)
    for col in 'ABDFJ':
        put(ws, f'{col}{r}', None, fill_=FILL_TOT)
    put(ws, 'A36', f'=IF(COUNT({keys})>{ST_N - 1},"⚠ 这段时间往来超过 {ST_N - 1} 笔，只列了前 {ST_N - 1} 笔，请把日期范围缩短","")',
        F_RED, align=AL, border=False)
    ws.merge_cells('A36:J36')
    ws.column_dimensions['K'].hidden = True
    widths(ws, {'A': 17, 'B': 13, 'C': 16, 'D': 14, 'E': 11, 'F': 13, 'G': 15, 'H': 14, 'I': 14, 'J': 20})
    ws.conditional_formatting.add(f'A{ST_DET0 + 1}:J{ST_DET1}',
                                  FormulaRule(formula=[f'AND($A{ST_DET0 + 1}<>"",MOD($A{ST_DET0 + 1},2)=0)'], fill=fill('FFF2F2F2')))
    ws.conditional_formatting.add(f'C{ST_DET0 + 1}:C{ST_DET1}',
                                  FormulaRule(formula=[f'LEFT($C{ST_DET0 + 1},2)="收款"'], font=Font(name=YH, sz=10, bold=True, color='FF00B050')))
    ws.print_title_rows = '38:38'
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ───────────────────────── 供应商对账单 ─────────────────────────
def build_sst(wb, ctx):
    ws = wb.create_sheet(SH_SST)
    ws.sheet_properties.tabColor = C_STMT[2:]
    title(ws, '供应商对账单', 'J', C_STMT,
          '💡 选供应商、起止日期，自动出和厂家的对账单：进了多少货、付了多少钱、还欠多少。数据来自【采购进货】【资金台帐】。')
    _sel_bar(ws, '🔍 供应商：', '供应商列表', ctx)
    sel, d0, d1 = '$B$4', '$G$4', '$I$4'
    ap0 = base_rng(S_AP0, SUP_R0, SUP_R1)
    opening = (f'SUMIF({SUP_NAMES},{sel},{ap0})+SUMIFS({buy(B_AMT)},{buy(B_SUP)},{sel},{buy(B_DATE)},">="&年初日,{buy(B_DATE)},"<"&{d0})'
               f'+SUMIFS({cash(K_NET)},{cash(K_SUP)},{sel},{cash(K_TO)},"{TO_AP}",{cash(K_DATE)},">="&年初日,{cash(K_DATE)},"<"&{d0})')
    _section(ws, 7, '💰 货款汇总', 'FF7030A0')
    heads = ['期初应付', '本期进货金额', '本期付款', '期末应付', '状态', '本期进货数量(件)', '本期厂家返利(已收)']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}8', t, F_HDR, fill('FF7030A0'), align=ACW)
    ws.row_dimensions[8].height = 30
    vals = [f'={opening}',
            f'=SUMIFS({buy(B_AMT)},{buy(B_SUP)},{sel},{buy(B_DATE)},">="&{d0},{buy(B_DATE)},"<"&({d1}+1))',
            f'=-SUMIFS({cash(K_NET)},{cash(K_SUP)},{sel},{cash(K_TO)},"{TO_AP}",{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&({d1}+1))',
            '=A9+B9-C9',
            '=IF(ROUND(D9,2)=0,"已结清",IF(D9>0,"我方欠款","我方预付"))',
            f'=SUMIFS({buy(B_QTY)},{buy(B_SUP)},{sel},{buy(B_DATE)},">="&{d0},{buy(B_DATE)},"<"&({d1}+1))'
            f'+SUMIFS({buy(B_GIFT)},{buy(B_SUP)},{sel},{buy(B_DATE)},">="&{d0},{buy(B_DATE)},"<"&({d1}+1))',
            f'=SUMIFS({cash(K_NET)},{cash(K_SUP)},{sel},{cash(K_TO)},"{TO_REB}",{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&({d1}+1))']
    for i, v in enumerate(vals):
        put(ws, f'{CL(i + 1)}9', v, Font(name=YH, sz=11, bold=True), fill('FFE4DFEC'),
            QTY if i == 5 else (None if i == 4 else MONEY2), AC)
    ws.row_dimensions[9].height = 24
    put(ws, 'A10', '说明：应付为正＝我们还欠厂家；为负＝我们预付了货款。进货数量含赠品。', F_NOTE, align=AL, border=False)
    ws.merge_cells('A10:J10')
    put(ws, 'A11', '以上往来核对无误。供应商签字（盖章）：', F_TXTB, align=AL, border=False)
    ws.merge_cells('A11:E11')
    put(ws, 'G11', '日期：', F_TXTB, align=AR_, border=False)
    ws.row_dimensions[11].height = 30

    _section(ws, 37, '📋 往来明细（按日期排）', 'FF4472C4')
    heads = ['序号', '日期', '摘要', '商品品类', '数量(件)', '进价', '进货金额', '付款金额', '应付余额', '备注']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}38', t, F_HDR, fill('FF4472C4'), align=ACW)
    ws.row_dimensions[38].height = 28
    keys = f'{SH_AUX}!$H${AUX_S0}:$H${AUX_S1}'
    b = lambda col: f'{SH_BUY}!${col}${BUY_R0}:${col}${BUY_R1}'
    c = lambda col: f'{SH_CASH}!${col}${CASH_R0}:${col}${CASH_R1}'
    r = ST_DET0
    put(ws, f'B{r}', f'={d0}', F_TXT, fmt=DATE, align=AC)
    put(ws, f'C{r}', '期初应付', F_TXTB, align=AC)
    put(ws, f'I{r}', '=A9', F_TXTB, fmt=MONEY2, align=AR_)
    for col in 'ADEFGHJ':
        put(ws, f'{col}{r}', None)
    for k in range(1, ST_N):
        r = ST_DET0 + k
        K = f'$K{r}'
        isb = f'{K}<={N_BUY}'
        ci = f'{K}-{N_BUY}'
        put(ws, f'K{r}', f'=IF({k}>{SH_AUX}!$H$2,"",IFERROR(MATCH(SMALL({keys},{k}),{keys},0),""))', F_NOTE, border=False)
        put(ws, f'A{r}', f'=IF({K}="","",{k})', F_TXT, align=AC)
        put(ws, f'B{r}', f'=IF({K}="","",IF({isb},INDEX({b(B_DATE)},{K}),INDEX({c(K_DATE)},{ci})))', F_TXT, fmt=DATE, align=AC)
        put(ws, f'C{r}', (f'=IF({K}="","",IF({isb},IF(N(INDEX({b(B_QTY)},{K}))<0,"退货","进货"),'
                          f'IF(INDEX({c(K_NET)},{ci})<=0,"付款","退款")&"（"&INDEX({c(K_ACC)},{ci})&"）"))'), F_TXT, align=AC)
        put(ws, f'D{r}', f'=IF({K}="","",IF({isb},INDEX({b(B_GOODS)},{K})&"",""))', F_TXT, align=AC)
        put(ws, f'E{r}', f'=IF({K}="","",IF({isb},N(INDEX({b(B_QTY)},{K}))+N(INDEX({b(B_GIFT)},{K})),""))', F_TXT, fmt=QTY, align=AC)
        put(ws, f'F{r}', f'=IF({K}="","",IF({isb},IF(INDEX({b(B_PRICE)},{K})="","",INDEX({b(B_PRICE)},{K})),""))', F_TXT, fmt=PRICE, align=AC)
        put(ws, f'G{r}', f'=IF({K}="","",IF({isb},N(INDEX({b(B_AMT)},{K})),""))', F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'H{r}', f'=IF({K}="","",IF({isb},"",-INDEX({c(K_NET)},{ci})))', F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'I{r}', f'=IF({K}="","",I{r - 1}+N(G{r})-N(H{r}))', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'J{r}', (f'=IF({K}="","",IF({isb},IF(N(INDEX({b(B_GIFT)},{K}))<>0,"含赠品"&INDEX({b(B_GIFT)},{K})&"件 ","")'
                          f'&INDEX({b(B_NOTE)},{K}),INDEX({c(K_MEMO)},{ci})&""))'), F_NOTE, align=AL)
    r = ST_TOT
    put(ws, f'C{r}', '本期合计', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'E{r}', f'=SUM(E{ST_DET0}:E{ST_DET1})', F_TXTB, FILL_TOT, QTY, AC)
    put(ws, f'G{r}', f'=SUM(G{ST_DET0}:G{ST_DET1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'H{r}', f'=SUM(H{ST_DET0}:H{ST_DET1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'I{r}', '=D9', F_TXTB, FILL_TOT, MONEY2, AR_)
    for col in 'ABDFJ':
        put(ws, f'{col}{r}', None, fill_=FILL_TOT)
    put(ws, 'A36', f'=IF(COUNT({keys})>{ST_N - 1},"⚠ 这段时间往来超过 {ST_N - 1} 笔，只列了前 {ST_N - 1} 笔，请把日期范围缩短","")',
        F_RED, align=AL, border=False)
    ws.merge_cells('A36:J36')
    # 供应商对账单没有「存条」块，13~35 行放一段本期按商品的进货小结
    _section(ws, 13, '📦 本期进货按商品小结', 'FF5B9BD5')
    heads = ['商品品类', '进货数量(件)', '赠品数量(件)', '进货金额', '平均进价', '', '']
    for i, t in enumerate(heads[:5]):
        put(ws, f'{CL(i + 1)}14', t, F_HDR, fill('FF5B9BD5'), align=ACW)
    gnames = GOODS_NAMES
    # 只列这个供应商在期间内进过的商品：用商品档案顺序，进货为 0 的显示空
    for k in range(1, 21):
        r = 14 + k
        # 第 k 个「这个供应商本期有进货」的商品在商品档案里的位置，存在 L 列
        put(ws, f'L{r}', f'=IFERROR(SMALL($M$15:$M${15 + BASE_R1 - BASE_R0},{k}),"")', F_NOTE, border=False)
    for i in range(BASE_R1 - BASE_R0 + 1):
        rr = 15 + i
        g = f'{SH_BASE}!${G_NAME}${BASE_R0 + i}'
        put(ws, f'M{rr}', (f'=IF({g}="","",IF(SUMIFS({buy(B_QTY)},{buy(B_SUP)},{sel},{buy(B_GOODS)},{g},{buy(B_DATE)},">="&{d0},'
                           f'{buy(B_DATE)},"<"&({d1}+1))+SUMIFS({buy(B_GIFT)},{buy(B_SUP)},{sel},{buy(B_GOODS)},{g},{buy(B_DATE)},">="&{d0},'
                           f'{buy(B_DATE)},"<"&({d1}+1))<>0,{i + 1},""))'), F_NOTE, border=False)
    for k in range(1, 21):
        r = 14 + k
        L = f'$L{r}'
        g = f'INDEX({gnames},{L})'
        cond = f'{buy(B_SUP)},{sel},{buy(B_GOODS)},{g},{buy(B_DATE)},">="&{d0},{buy(B_DATE)},"<"&({d1}+1)'
        put(ws, f'A{r}', f'=IF({L}="","",{g})', F_TXT, align=AC)
        put(ws, f'B{r}', f'=IF({L}="","",SUMIFS({buy(B_QTY)},{cond}))', F_TXT, fmt=QTY, align=AC)
        put(ws, f'C{r}', f'=IF({L}="","",SUMIFS({buy(B_GIFT)},{cond}))', F_TXT, fmt=QTY, align=AC)
        put(ws, f'D{r}', f'=IF({L}="","",SUMIFS({buy(B_AMT)},{cond}))', F_TXT, fmt=MONEY2, align=AC)
        put(ws, f'E{r}', f'=IF(OR({L}="",N(B{r})+N(C{r})=0),"",D{r}/(B{r}+C{r}))', F_TXT, fmt=PRICE, align=AC)
    for col in 'KLM':
        ws.column_dimensions[col].hidden = True
    widths(ws, {'A': 17, 'B': 13, 'C': 16, 'D': 14, 'E': 11, 'F': 13, 'G': 15, 'H': 14, 'I': 14, 'J': 22})
    ws.conditional_formatting.add(f'A{ST_DET0 + 1}:J{ST_DET1}',
                                  FormulaRule(formula=[f'AND($A{ST_DET0 + 1}<>"",MOD($A{ST_DET0 + 1},2)=0)'], fill=fill('FFF2F2F2')))
    ws.print_title_rows = '38:38'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


# ───────────────────────── 收付款汇总 ─────────────────────────
RPS_M0 = 7                     # 1 月在第 7 行
RPS_CUS0 = 24                  # 客户块第一行
RPS_CUS1 = RPS_CUS0 + (OV_R1 - OV_R0)
RPS_SUP0 = RPS_CUS0
RPS_SUP1 = RPS_SUP0 + (SUP_R1 - SUP_R0)


def build_rps(wb, ctx):
    ws = wb.create_sheet(SH_RPS)
    ws.sheet_properties.tabColor = C_RPT[2:]
    title(ws, '收付款汇总', 'Q', C_RPT,
          '💡 全自动：上面是本年度逐月的收付款（从【资金台帐】按收支项目分类汇总）；下面按客户、按供应商汇总「起止日期」内的往来，'
          '期末欠款＝期初＋本期领用/进货－本期收/付款。要看某一家的逐笔明细，去【客户对账单】【供应商对账单】。')
    # 按月
    _section(ws, 5, '📅 按月收付款汇总（本年度，不含内部转账）', C_CASH, c1='O')
    heads = ['月份', '客户回款', '付供应商货款', '费用支出', '报销付款', '厂家返利', '其他收入', '股东投入', '股东提取',
             '借款净额\n(借入－归还)', '其他往来净额\n(收－付)', '收入合计', '支出合计', '本月净额', '月末资金余额']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}6', t, F_HDR, fill(C_CASH), align=ACW)
    ws.row_dimensions[6].height = 34
    acc0 = f'SUM({SH_BASE}!${A_BAL0}${ACC_R0}:${A_BAL0}${ACC_R1})'
    for m in range(1, 13):
        r = RPS_M0 + m - 1
        mc = f'{cash(K_DATE)},">="&DATE(年度,{m},1),{cash(K_DATE)},"<"&DATE(年度,{m}+1,1)'
        s = lambda to, sign='': f'={sign}SUMIFS({cash(K_NET)},{cash(K_TO)},"{to}",{mc})'
        put(ws, f'A{r}', m, F_TXTB, FILL_SUBH, '0"月"', AC)
        cols = [(TO_AR, ''), (TO_AP, '-'), (TO_EXP, '-'), (TO_REIMB, '-'), (TO_REB, ''), (TO_OI, ''), (TO_INV, ''),
                (TO_DRAW, '-'), (TO_LOAN, ''), (TO_OTH, '')]
        for j, (to, sign) in enumerate(cols):
            put(ws, f'{CL(2 + j)}{r}', s(to, sign), F_TXT, fmt=MONEY2, align=AR_)
        ci, co, cn, cb = CL(2 + len(cols)), CL(3 + len(cols)), CL(4 + len(cols)), CL(5 + len(cols))
        put(ws, f'{ci}{r}', f'=SUMIFS({cash(K_IN)},{mc},{cash(K_TO)},"<>{TO_XFER}")', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'{co}{r}', f'=SUMIFS({cash(K_OUT)},{mc},{cash(K_TO)},"<>{TO_XFER}")', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'{cn}{r}', f'={ci}{r}-{co}{r}', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'{cb}{r}', f'={acc0}+SUMIFS({cash(K_NET)},{cash(K_DATE)},">="&年初日,{cash(K_DATE)},"<"&DATE(年度,{m}+1,1))',
            F_TXTB, fmt=MONEY2, align=AR_)
    r = RPS_M0 + 12
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for col in 'BCDEFGHIJKLMN':
        put(ws, f'{col}{r}', f'=SUM({col}{RPS_M0}:{col}{RPS_M0 + 11})', F_TXTB, FILL_TOT, MONEY2, AR_)
    put(ws, f'O{r}', f'=O{RPS_M0 + 11}', F_TXTB, FILL_TOT, MONEY2, AR_)

    # 起止日期
    put(ws, 'A21', '起始日期：', F_KPI_L, align=AR_, border=False)
    put(ws, 'B21', '=年初日', Font(name=YH, sz=11, bold=True), FILL_SEL, DATE, AC)
    put(ws, 'C21', '截止日期：', F_KPI_L, align=AR_, border=False)
    put(ws, 'D21', '=年末日', Font(name=YH, sz=11, bold=True), FILL_SEL, DATE, AC)
    put(ws, 'E21', '← 改这两格看任意一段时间（下面两块跟着变）', F_NOTE, align=AL, border=False)
    ws.merge_cells('E21:H21')
    d0, d1 = '$B$21', '$D$21'
    # 客户块 A:H
    ws.merge_cells('A22:H22')
    put(ws, 'A22', '👥 客户收款汇总（按总览汇总的客户顺序）', Font(name=YH, sz=12, bold=True, color='FFFFFFFF'), fill(C_OV), align=AL)
    heads = ['序号', '客户', '期初欠款', '本期领用', '本期收款', '期末欠款', '状态', '最近收款日']
    for i, t in enumerate(heads):
        put(ws, f'{CL(i + 1)}23', t, F_HDR, fill(C_OV), align=ACW)
    for i in range(OV_R1 - OV_R0 + 1):
        r = RPS_CUS0 + i
        vr = OV_R0 + i
        b = f'$B{r}'
        put(ws, f'A{r}', f'=IF(B{r}="","",{i + 1})', F_AUTO, align=AC)
        put(ws, f'B{r}', f'=IF({at(CUS_NAMES, i + 1)}="","",{at(CUS_NAMES, i + 1)})', F_TXT, align=AC)
        put(ws, f'C{r}', (f'=IF({b}="","",N({at(CUS_OPEN, i + 1)})+SUMIFS({out(O_AMT)},{out(O_CUS)},{b},{out(O_DATE)},">="&年初日,{out(O_DATE)},"<"&{d0})'
                          f'-SUMIFS({cash(K_NET)},{cash(K_CUS)},{b},{cash(K_TO)},"{TO_AR}",{cash(K_DATE)},">="&年初日,{cash(K_DATE)},"<"&{d0}))'),
            F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'D{r}', f'=IF({b}="","",SUMIFS({out(O_AMT)},{out(O_CUS)},{b},{out(O_DATE)},">="&{d0},{out(O_DATE)},"<"&({d1}+1)))',
            F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'E{r}', (f'=IF({b}="","",SUMIFS({cash(K_NET)},{cash(K_CUS)},{b},{cash(K_TO)},"{TO_AR}",'
                          f'{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&({d1}+1)))'), F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'F{r}', f'=IF({b}="","",C{r}+D{r}-E{r})', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'G{r}', f'=IF({b}="","",IF(ROUND(F{r},2)=0,"已结清",IF(F{r}>0,"有欠款","预存款")))', F_TXT, align=AC)
        put(ws, f'H{r}', (f'=IF({b}="","",IFERROR(_xlfn.AGGREGATE(14,6,{cash(K_DATE)}/(({cash(K_CUS)}={b})*({cash(K_TO)}="{TO_AR}")'
                          f'*({cash(K_NET)}>0)*({cash(K_DATE)}<{d1}+1)),1),0))'), F_TXT,
            fmt='yyyy/mm/dd;;', align=AC)
    r = RPS_CUS1 + 1
    put(ws, f'B{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for col in 'CDEF':
        put(ws, f'{col}{r}', f'=SUM({col}{RPS_CUS0}:{col}{RPS_CUS1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    ws.conditional_formatting.add(f'F{RPS_CUS0}:F{RPS_CUS1}', FormulaRule(formula=[f'AND(ISNUMBER(F{RPS_CUS0}),F{RPS_CUS0}>0.005)'],
                                                                          font=Font(name=YH, sz=10, bold=True, color='FFC00000')))
    # 供应商块 J:Q
    ws.merge_cells('J22:Q22')
    put(ws, 'J22', '🏭 供应商付款汇总', Font(name=YH, sz=12, bold=True, color='FFFFFFFF'), fill(C_BUY), align=AL)
    heads = ['序号', '供应商', '期初应付', '本期进货', '本期付款', '期末应付', '状态', '最近付款日']
    for i, t in enumerate(heads):
        put(ws, f'{CL(10 + i)}23', t, F_HDR, fill(C_BUY), align=ACW)
    ws.row_dimensions[23].height = 30
    for i in range(SUP_R1 - SUP_R0 + 1):
        r = RPS_SUP0 + i
        sr = SUP_R0 + i
        k = f'$K{r}'
        put(ws, f'J{r}', f'=IF(K{r}="","",{i + 1})', F_AUTO, align=AC)
        put(ws, f'K{r}', f'=IF({at(SUP_NAMES, i + 1)}="","",{at(SUP_NAMES, i + 1)})', F_TXT, align=AC)
        put(ws, f'L{r}', (f'=IF({k}="","",N({at(base_rng(S_AP0, SUP_R0, SUP_R1), i + 1)})+SUMIFS({buy(B_AMT)},{buy(B_SUP)},{k},{buy(B_DATE)},">="&年初日,{buy(B_DATE)},"<"&{d0})'
                          f'+SUMIFS({cash(K_NET)},{cash(K_SUP)},{k},{cash(K_TO)},"{TO_AP}",{cash(K_DATE)},">="&年初日,{cash(K_DATE)},"<"&{d0}))'),
            F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'M{r}', f'=IF({k}="","",SUMIFS({buy(B_AMT)},{buy(B_SUP)},{k},{buy(B_DATE)},">="&{d0},{buy(B_DATE)},"<"&({d1}+1)))',
            F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'N{r}', (f'=IF({k}="","",-SUMIFS({cash(K_NET)},{cash(K_SUP)},{k},{cash(K_TO)},"{TO_AP}",'
                          f'{cash(K_DATE)},">="&{d0},{cash(K_DATE)},"<"&({d1}+1)))'), F_TXT, fmt=MONEY2, align=AR_)
        put(ws, f'O{r}', f'=IF({k}="","",L{r}+M{r}-N{r})', F_TXTB, fmt=MONEY2, align=AR_)
        put(ws, f'P{r}', f'=IF({k}="","",IF(ROUND(O{r},2)=0,"已结清",IF(O{r}>0,"我方欠款","我方预付")))', F_TXT, align=AC)
        put(ws, f'Q{r}', (f'=IF({k}="","",IFERROR(_xlfn.AGGREGATE(14,6,{cash(K_DATE)}/(({cash(K_SUP)}={k})*({cash(K_TO)}="{TO_AP}")'
                          f'*({cash(K_NET)}<0)*({cash(K_DATE)}<{d1}+1)),1),0))'), F_TXT,
            fmt='yyyy/mm/dd;;', align=AC)
    r = RPS_SUP1 + 1
    put(ws, f'K{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for col in 'LMNO':
        put(ws, f'{col}{r}', f'=SUM({col}{RPS_SUP0}:{col}{RPS_SUP1})', F_TXTB, FILL_TOT, MONEY2, AR_)
    widths(ws, {'A': 10, 'B': 16, 'C': 13, 'D': 13, 'E': 13, 'F': 13, 'G': 12, 'H': 12, 'I': 12,
                'J': 12, 'K': 14, 'L': 13, 'M': 13, 'N': 13, 'O': 13, 'P': 11, 'Q': 11})
    ws.freeze_panes = 'A7'
    return ws
