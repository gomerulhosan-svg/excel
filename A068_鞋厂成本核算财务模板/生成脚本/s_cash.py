# -*- coding: utf-8 -*-
"""【资金日记账】（4 个资金账户混合录入、逐行即时余额）和【手工分录】（一借一贷，调账 / 暂估 / 计提用）。
   ★ 隐藏接口列（R～AG / N～U）的含义见 layout.py，这里一字不差地实现；AH～AL / V～W 是本表自己用的帮手列（也隐藏）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from common import *
from layout import *          # 注意：layout 的 C_IN（录入组颜色）覆盖 common 的同名浅绿色

HDR_IN, HDR_AUTO, HDR_CHK = 'FF2F75B5', 'FF7F7F7F', 'FF833C0C'
KPI_FILL = fill('FFD9E1F2')
F_VAL = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_CNT = Font(name=YH, sz=11, bold=True, color='FFC00000')
F_GREEN = Font(name=YH, sz=10, bold=True, color='FF00B050')
F_REDB = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_REDN = Font(name=YH, sz=10, color='FFC00000')
FILL_YEL = fill('FFFFEB9C')
SHRINK = Alignment(horizontal='center', vertical='center', shrink_to_fit=True)
CNT = '0" 笔"'

# 收支类别里几个「自动认」要用到的名字（跟 layout.CATS 一致）
_CAT = {c[0]: c for c in CATS}
for _n in ('收货款', '付材料款', '付加工费', '费用支出', '内部转账'):
    assert _n in _CAT, _n
XFER = '内部转账'
# 对方科目 → 成本组件（开发说明：400101 材料 / 400102 外发 / 400103 人工 / 4101 制造）
_COMP_OF = [(COMP_CODE['材料'], '材料'), (COMP_CODE['外发'], '外发'), (COMP_CODE['人工'], '人工'), (MOH, '制造')]
# 系统月末自动结转、手工分录不能用的科目
_SYS_CODES = [COMP_CODE['制造']] + [COGS_CODE[c] for c in COMPS]
# 费用归类 → 现金流量项目：跟默认「支付其他与经营有关的现金」不一样的归类（税金及附加、所得税）
_CF_OTHER = '支付其他与经营有关的现金'
_CF_SPECIAL = {}
for _k, (_code, _cf) in FE_CLASSES.items():
    if _cf != _CF_OTHER:
        _CF_SPECIAL.setdefault(_cf, []).append(_k)


def comp_of(x):
    """科目编码（文本）→ 成本组件，不是成本科目 → "" """
    f = '""'
    for code, c in reversed(_COMP_OF):
        f = f'IF({x}="{code}","{c}",{f})'
    return f


def cf_of_fee(fc):
    """费用归类 → 现金流量项目（税金及附加、所得税 → 支付的各项税费；其余 → 支付其他与经营有关的现金）"""
    f = f'"{_CF_OTHER}"'
    for cf, classes in _CF_SPECIAL.items():
        cond = 'OR(' + ','.join(f'{fc}="{k}"' for k in classes) + ')'
        f = f'IF({cond},"{cf}",{f})'
    return f


def chain(pairs, last):
    """按顺序判断，命中第一个就停：[(条件, 结果), ...]，都不中 → last"""
    f = last
    for cond, res in reversed(pairs):
        f = f'IF({cond},{res},{f})'
    return f


def _kpi(ws, coord_l, label, coord_v, formula, fmt, font=F_VAL):
    put(ws, coord_l, label, F_KPI_L, KPI_FILL, align=ACW)
    put(ws, coord_v, formula, font, KPI_FILL, fmt, AC)


def _check_fmt(ws, rg, first):
    """校验列：✗ 红底红字、⚠ 黄底、√ 绿字"""
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="✗"'], fill=FILL_WARN, font=F_REDB))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], fill=FILL_YEL))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="√"'], font=F_GREEN))


# ═══════════════════════════ 资金日记账 ═══════════════════════════
# 本表自己的帮手列（隐藏，在 ★ 接口列 AG 的右边）
X_CTI, X_NEED, X_ACI, X_FEI, X_TOI = 'AH', 'AI', 'AJ', 'AK', 'AL'


def build_cash(wb, ctx):
    ws = wb[SH_CASH]
    assert (CI(J_OUT) - CI(J_DATE), CI(J_NOTE) - CI(J_TOACC)) == (7, 3)
    widths(ws, {J_SEQ: 8, J_DATE: 11, J_ACC: 9, J_CAT: 11, J_UNIT: 16, J_FEE: 11, J_MEMO: 24, J_IN: 12, J_OUT: 12,
                J_BAL: 13, J_TBAL: 13, J_TOACC: 10, J_STY: 10, J_NO: 11, J_NOTE: 16, J_CHK: 40, 'Q': 2})
    title(ws, '资 金 日 记 账（银行 · 现金 · 支付宝 混在一张表 · 逐行即时余额）', J_CHK, C_IN,
          '💡 所有账户（银行1、银行2、现金、支付宝……）混在这一张表里，按时间顺序一笔一行：日期、账户、金额必填，收进来的填「收入」，付出去的填「支出」，只填一边。'
          '「收支类别」可以不选：往来单位选了材料供应商自动算「付材料款」、选了外发加工厂算「付加工费」、选了客户算「收货款」，选了费用项目自动算「费用支出」。'
          '自己账户之间转钱（银行取现、银行转支付宝）只记一行：从哪个账户转出就在那个账户记「支出」，「对方账户」选转进去的账户，两个账户的余额都会跟着变。'
          '发工资选「付工资」；没有送货单、现场买的零星材料选「现付材料」；某个款专用的东西（模具、专用材料）在「用在哪」选款式，成本直接记到这个款。'
          '最右边「校验」：✗ 的行进不了账（报表里没有），照提示改；⚠ 是提醒，照样记账。记错的行清空内容就行，不要删整行、不要插行（右边有隐藏公式）。')

    # ── 第 3 行：各资金账户当前余额（名称 / 余额两格一对，对应【基础资料】② 第 1～8 个账户）
    n_ac = AC_R1 - AC_R0 + 1
    assert 2 * n_ac <= CI(J_CHK)
    now_r = rng(SH_BASE, AC_NOW, AC_R0, AC_R1)
    for k in range(1, n_ac + 1):
        cn, cv = CL(2 * k - 1), CL(2 * k)
        put(ws, f'{cn}3', f'=IF(INDEX({AC_NAMES},{k})&""="","",INDEX({AC_NAMES},{k})&"")', F_KPI_L, KPI_FILL, align=SHRINK)
        put(ws, f'{cv}3', f'=IF({cn}3="","",N(INDEX({now_r},{k})))', F_VAL, KPI_FILL, MONEY, AC)
    ws.row_dimensions[3].height = 24
    # ── 第 4 行：总余额、本年收入、本年支出、笔数、✗、⚠、有金额没进账
    _kpi(ws, 'A4', '总余额', 'B4', f'=SUM({now_r})', MONEY)
    _kpi(ws, 'C4', '本年收入', 'D4', f'=SUMIFS({jr(J_NET)},{jr(J_OK)},1,{jr(J_CATX)},"<>{XFER}",{jr(J_NET)},">0")', MONEY)
    _kpi(ws, 'E4', '本年支出', 'F4', f'=-SUMIFS({jr(J_NET)},{jr(J_OK)},1,{jr(J_CATX)},"<>{XFER}",{jr(J_NET)},"<0")', MONEY)
    _kpi(ws, 'G4', '笔数', 'H4', f'=COUNT({jr(J_SEQ)})', CNT)
    _kpi(ws, 'I4', '✗ 待改', 'J4', f'=COUNTIF({jr(J_CHK)},"✗*")', CNT, F_CNT)
    _kpi(ws, 'K4', '⚠ 提醒', 'L4', f'=COUNTIF({jr(J_CHK)},"⚠*")', CNT, F_CNT)
    _kpi(ws, 'M4', '有金额\n没进账', 'N4', f'=COUNTIFS({jr(J_NET)},"<>0",{jr(J_OK)},0)', CNT, F_CNT)
    put(ws, 'O4', '收入、支出不含内部转账；「没进账」＝有金额但校验是 ✗ 的笔数（改好才进报表）', F_NOTE, KPI_FILL, align=ALW)
    ws.merge_cells('O4:P4')
    ws.row_dimensions[4].height = 30
    for c in ('J4', 'N4'):
        ws.conditional_formatting.add(c, FormulaRule(formula=[f'N({c})>0'], fill=FILL_WARN))
    ws.conditional_formatting.add('L4', FormulaRule(formula=['N(L4)>0'], fill=FILL_YEL))

    # ── 表头
    heads = [(J_SEQ, '序号'), (J_DATE, '日期'), (J_ACC, '账户'), (J_CAT, '收支类别\n（可不选）'), (J_UNIT, '往来单位 / 人'),
             (J_FEE, '费用项目'), (J_MEMO, '摘要'), (J_IN, '收入'), (J_OUT, '支出'), (J_BAL, '本账户余额'), (J_TBAL, '总余额'),
             (J_TOACC, '对方账户\n（内部转账）'), (J_STY, '用在哪\n（款式）'), (J_NO, '单号 / 票号'), (J_NOTE, '备注'), (J_CHK, '校验')]
    auto_cols = (J_SEQ, J_BAL, J_TBAL)
    for col, t in heads:
        c = HDR_CHK if col == J_CHK else (HDR_AUTO if col in auto_cols else HDR_IN)
        put(ws, f'{col}{J_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[J_HDR].height = 34
    hidden = [(J_NET, 'NET 净额'), (J_CM, 'CM 月份'), (J_CATX, 'CATX 实际类别'), (J_ACODE, 'ACODE 本账户科目'), (J_OPP, 'OPP 对方科目'),
              (J_DR, 'DR 借方'), (J_CR, 'CR 贷方'), (J_AMT, 'AMT 金额'), (J_UTYPE, 'UTYPE 单位类型'), (J_COMP, 'COMP 成本组件'),
              (J_CAMT, 'CAMT 成本额'), (J_CKEY, 'CKEY 成本键'), (J_CF, 'CF 现金流量项目'), (J_TOX, 'TOX 对方账户'), (J_OK, 'OK 能记账'),
              (J_FCLS, 'FCLS 费用归类'), (X_CTI, '类别序号'), (X_NEED, '还要选'), (X_ACI, '账户序号'), (X_FEI, '费用项目序号'),
              (X_TOI, '对方账户序号')]
    for col, t in hidden:
        ws[f'{col}{J_HDR}'] = t
        ws[f'{col}{J_HDR}'].font = F_HELP

    yr, op = P['YEAR'], P['OPEN']
    r0 = J_R0
    for r in range(J_R0, J_R1 + 1):
        A, B, C, D, E, F, G, H, I, J, K, L, M = (f'{c}{r}' for c in (J_SEQ, J_DATE, J_ACC, J_CAT, J_UNIT, J_FEE, J_MEMO, J_IN,
                                                                        J_OUT, J_BAL, J_TBAL, J_TOACC, J_STY))
        Pc = f'{J_CHK}{r}'
        R, S, T, U, V, Z, AA, AF, FC = (f'{c}{r}' for c in (J_NET, J_CM, J_CATX, J_ACODE, J_OPP, J_UTYPE, J_COMP, J_OK, J_FCLS))
        CTI, NEED, ACI, FEI, TOI = (f'{c}{r}' for c in (X_CTI, X_NEED, X_ACI, X_FEI, X_TOI))
        f = {}
        f[J_SEQ] = f'=IF(COUNTA({J_DATE}{r}:{J_OUT}{r},{J_TOACC}{r}:{J_NOTE}{r})=0,"",ROW()-{J_HDR})'
        f[J_NET] = f'=ROUND(N({H})-N({I}),2)'
        f[J_CM] = '=' + ym_month(B, yr)
        f[X_ACI] = f'=IF({C}="",0,IFERROR(MATCH({C}&"",{AC_NAMES},0),0))'
        f[X_TOI] = f'=IF({L}="",0,IFERROR(MATCH({L}&"",{AC_NAMES},0),0))'
        f[X_FEI] = f'=IF(TRIM({F}&"")="",0,IFERROR(MATCH(TRIM({F}&""),{FE_NAMES},0),0))'
        f[J_UTYPE] = f'=IF(TRIM({E}&"")="","",IFERROR(INDEX({UN_TYPES_R},MATCH(TRIM({E}&""),{UN_NAMES},0))&"",""))'
        f[J_CATX] = (f'=IF(TRIM({D}&"")<>"",TRIM({D}&""),IF({R}=0,"",IF({L}<>"","{XFER}",IF(TRIM({F}&"")<>"","费用支出",'
                     f'IF({Z}="材料供应商","付材料款",IF({Z}="外发加工厂","付加工费",IF({Z}="客户","收货款","")))))))')
        f[X_CTI] = f'=IF({T}="",0,IFERROR(MATCH({T},{CT_NAMES},0),0))'
        f[X_NEED] = f'=IF({CTI}=0,"",TRIM(INDEX({CT_NEEDS},{CTI})&""))'
        f[J_ACODE] = f'=IF({ACI}=0,"",TRIM(INDEX({AC_CODES},{ACI})&""))'
        f[J_OPP] = (f'=IF({CTI}=0,"",IF({NEED}="费用项目",IF({FEI}=0,"",TRIM(INDEX({FE_CODES},{FEI})&"")),'
                    f'IF({NEED}="对方账户",IF({TOI}=0,"",TRIM(INDEX({AC_CODES},{TOI})&"")),TRIM(INDEX({CT_OPPS},{CTI})&""))))')
        f[J_DR] = f'=IF({R}>0,{U},IF({R}<0,{V},""))'
        f[J_CR] = f'=IF({R}>0,{V},IF({R}<0,{U},""))'
        f[J_AMT] = f'=ABS({R})'
        f[J_COMP] = '=' + comp_of(V)
        f[J_CAMT] = f'=IF(AND({AA}<>"",{AF}=1),-{R},0)'
        m = f'TRIM({M}&"")'
        f[J_CKEY] = (f'=IF(AND({AA}<>"",{AF}=1),IF({m}="","公||"&{AA},IF(COUNTIF({ST_KEYS},{m})>0,"款|"&{m}&"|"&{AA},'
                     f'"公||"&{AA})),"")')
        f[J_CF] = (f'=IF(OR({AF}<>1,{T}="{XFER}"),"",IF({NEED}="费用项目",{cf_of_fee(FC)},'
                   f'TRIM(INDEX({CT_CFS},{CTI})&"")))')
        f[J_TOX] = f'=IF({T}="{XFER}",{L}&"","")'
        f[J_OK] = f'=IF(AND({Pc}<>"",LEFT({Pc},1)<>"✗"),1,0)'
        f[J_FCLS] = f'=IF(AND({NEED}="费用项目",{FEI}>0),TRIM(INDEX({FE_CLSS},{FEI})&""),"")'
        # 本账户余额 / 总余额（逐行即时：从第一行到本行）
        f[J_BAL] = (f'=IF(OR({C}="",{R}=0),"",ROUND(IF({ACI}>0,N(INDEX({AC_OPENS},{ACI})),0)'
                    f'+SUMIFS(${J_NET}${r0}:${J_NET}{r},${J_ACC}${r0}:${J_ACC}{r},{C})'
                    f'-SUMIFS(${J_NET}${r0}:${J_NET}{r},${J_TOX}${r0}:${J_TOX}{r},{C}),2))')
        f[J_TBAL] = (f'=IF({R}=0,"",ROUND(SUM({AC_OPENS})'
                     f'+SUMIFS(${J_NET}${r0}:${J_NET}{r},${J_CATX}${r0}:${J_CATX}{r},"<>{XFER}"),2))')
        # 校验（按顺序，命中第一个就停）
        dup = (f'IF({T}="{XFER}",COUNTIFS({jr(J_ACC)},{L}&"",{jr(J_TOX)},{C}&"",{jr(J_NET)},-{R},{jr(J_DATE)},{B}),0)>0')
        chk = [
            (f'{C}=""', '"✗ 没选账户：在「账户」下拉里选"'),
            (f'{ACI}=0', '"✗ 账户不在【基础资料】②：从下拉选，或先到【基础资料】② 加这个账户"'),
            (f'{U}=""', '"✗ 这个账户在【基础资料】② 没选类型（算不出科目编码）：去选 银行 / 现金 / 支付宝微信"'),
            (f'NOT(ISNUMBER({B}))', '"✗ 日期要填成日期（如 2026/9/5）"'),
            (f'YEAR({B})<>{yr}', '"✗ 不在本会计年度（年份看【基础资料】）：别的年份的记到那一年的账本里"'),
            (f'{B}<{op}', '"✗ 早于建账日期（建账前的钱算在账户期初余额里，不用记）"'),
            (f'AND(N({H})<>0,N({I})<>0)', '"✗ 收入、支出只填一边（又收又付的拆成两行）"'),
            (f'OR(AND({H}<>"",NOT(ISNUMBER({H}))),AND({I}<>"",NOT(ISNUMBER({I}))))', '"✗ 金额要填数字（不要带「元」、空格、文字）"'),
            (f'OR(N({H})<0,N({I})<0)', '"✗ 金额不要填负数：退回来的钱记在收入栏、付出去的记在支出栏"'),
            (f'{R}=0', '"✗ 没有金额"'),
            (f'{T}=""', '"✗ 认不出收支类别：在「收支类别」选（或选往来单位 / 费用项目）"'),
            (f'{CTI}=0', '"✗ 收支类别不在【基础资料】③：从下拉选"'),
            (f'AND({NEED}="单位",TRIM({E}&"")="")', f'"✗ 这一类要选往来单位：「"&{T}&"」是跟哪一家 / 哪个人"'),
            (f'AND({NEED}="单位",{Z}="")', '"✗ 往来单位不在【往来单位】里：先去登记（名字要一模一样）"'),
            (f'AND({NEED}="费用项目",TRIM({F}&"")="")', '"✗ 费用支出要选费用项目（水电、房租、运费……）"'),
            (f'AND({NEED}="费用项目",{FEI}=0)', '"✗ 费用项目不在【基础资料】④：从下拉选，或先到 ④ 加上"'),
            (f'AND({NEED}="对方账户",{L}="")', '"✗ 内部转账要选对方账户（钱转进去的那个账户）"'),
            (f'AND({NEED}="对方账户",{TOI}=0)', '"✗ 对方账户不在资金账户里（【基础资料】②）：从下拉选"'),
            (f'AND({NEED}="对方账户",{TOI}={ACI})', '"✗ 对方账户不能是本账户"'),
            (f'{V}=""', '"✗ 对方科目找不到：看【基础资料】③ 这一类的「对方科目」（费用项目看 ④ 的「归类」）"'),
            (f'COUNTIF({COA_CODES},{V})=0', f'"✗ 对方科目 "&{V}&" 不在【会计科目表】：先去加科目，或改【基础资料】③"'),
            (dup, '"⚠ 对方账户那边好像也记了这笔转账（同一天、同金额）：转账只记一行，清空另一行"'),
            (f'AND({AA}<>"",{m}<>"",COUNTIF({ST_KEYS},{m})=0)', '"⚠ 用在哪的款式没在【款式档案】：先按公共分摊（登记了款式就自动记到这个款）"'),
            (f'AND({AA}="",{m}<>"")', '"⚠ 这一类不算成本，「用在哪」不起作用"'),
            (f'AND({T}="付材料款",{Z}<>"材料供应商")', f'"⚠ 付材料款的单位类型不对：【往来单位】里是「"&{Z}&"」，不是材料供应商"'),
            (f'AND({T}="付加工费",{Z}<>"外发加工厂")', f'"⚠ 付加工费的单位类型不对：【往来单位】里是「"&{Z}&"」，不是外发加工厂"'),
            (f'AND({T}="收货款",{Z}<>"客户")', f'"⚠ 收货款的单位不是客户：【往来单位】里是「"&{Z}&"」"'),
            (f'AND({L}<>"",{T}<>"{XFER}")', '"⚠ 选了对方账户但类别不是内部转账（对方账户不起作用）"'),
            (f'AND(ISNUMBER({J}),{J}<0)', '"⚠ 账户余额成负数了：是不是漏记了收入，或期初余额不对"'),
        ]
        f[J_CHK] = f'=IF({A}="","",' + chain(chk, f'IF(TRIM({D}&"")="","√ 自动认作："&{T},"√")') + ')'
        for col, v in f.items():
            ws[f'{col}{r}'] = v

    cols = [J_SEQ, J_DATE, J_ACC, J_CAT, J_UNIT, J_FEE, J_MEMO, J_IN, J_OUT, J_BAL, J_TBAL, J_TOACC, J_STY, J_NO, J_NOTE, J_CHK]
    in_cols = [c for c in cols if c not in auto_cols and c != J_CHK]
    style_rows(ws, J_R0, J_R1, cols, auto=list(auto_cols) + [J_CHK],
               fmts={J_SEQ: INT, J_DATE: DATE, J_IN: MONEY, J_OUT: MONEY, J_BAL: MONEY, J_TBAL: MONEY, J_STY: '@', J_NO: '@'},
               aligns={J_UNIT: AL, J_MEMO: AL, J_NOTE: AL, J_CHK: AL, J_IN: AR, J_OUT: AR, J_BAL: AR, J_TBAL: AR},
               fills={c: FILL_IN for c in in_cols})
    for r in range(J_R0, J_R1 + 1):
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hidden])

    # 下拉、验证
    dv_date(ws, f'{J_DATE}{J_R0}:{J_DATE}{J_R1}')
    dv_list(ws, f'{J_ACC}{J_R0}:{J_ACC}{J_R1}', f'={AC_NAMES}', '哪个账户的钱（【基础资料】② 登记的）')
    dv_list(ws, f'{J_CAT}{J_R0}:{J_CAT}{J_R1}', f'={CT_NAMES}',
            '可以不选：选了供应商自动算付材料款、选了客户算收货款、选了费用项目算费用支出、选了对方账户算内部转账')
    dv_list(ws, f'{J_UNIT}{J_R0}:{J_UNIT}{J_R1}', f'={UN_NAMES}', '供应商 / 加工厂 / 客户从下拉选；个人名字可以直接写', stop=False)
    dv_list(ws, f'{J_FEE}{J_R0}:{J_FEE}{J_R1}', f'={FE_NAMES}', '水电、房租、运费……（【基础资料】④）')
    dv_list(ws, f'{J_TOACC}{J_R0}:{J_TOACC}{J_R1}', f'={AC_NAMES}', '只有自己账户之间转钱才选：钱转进去的那个账户')
    dv_list(ws, f'{J_STY}{J_R0}:{J_STY}{J_R1}', f'={ST_CODES}', '某个款专用的才选（模具、专用材料）；空着＝公共，按交货双数分',
            stop=False)
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True,
                        errorStyle='warning', errorTitle='金额', error='填正数：退回来的钱记在收入栏、付出去的记在支出栏')
    ws.add_data_validation(dv)
    dv.add(f'{J_IN}{J_R0}:{J_OUT}{J_R1}')

    # 条件格式
    _check_fmt(ws, f'{J_CHK}{J_R0}:{J_CHK}{J_R1}', f'${J_CHK}{J_R0}')
    for c in (J_BAL, J_TBAL):
        ws.conditional_formatting.add(f'{c}{J_R0}:{c}{J_R1}', FormulaRule(
            formula=[f'AND(ISNUMBER(${c}{J_R0}),${c}{J_R0}<0)'], font=F_REDB))
    for k in range(1, n_ac + 1):
        c = f'{CL(2 * k)}3'
        ws.conditional_formatting.add(c, FormulaRule(formula=[f'AND(ISNUMBER({c}),{c}<0)'], font=F_REDB))
    ws.auto_filter.ref = f'{J_SEQ}{J_HDR}:{J_CHK}{J_R1}'
    ws.freeze_panes = f'{J_ACC}{J_R0}'

    # 演示数据
    for i, row in enumerate(ctx.get(SH_CASH, [])):
        r = J_R0 + i
        assert r <= J_R1
        for c, v in row.items():
            if v is not None:
                ws[f'{c}{r}'] = v
    return ws


# ═══════════════════════════ 手工分录 ═══════════════════════════
X_DRI, X_CRI = 'V', 'W'            # 帮手列：借方 / 贷方科目在【会计科目表】第几行（0＝没有）


def build_mj(wb, ctx):
    ws = wb[SH_MJ]
    widths(ws, {MJ_SEQ: 5, MJ_DATE: 11, MJ_MEMO: 26, MJ_DRIN: 10, MJ_DRN: 20, MJ_CRIN: 10, MJ_CRN: 20, MJ_AMTIN: 12,
                MJ_UNIT: 16, MJ_FEE: 11, MJ_NOTE: 18, MJ_CHK: 44, 'M': 2})
    title(ws, '手 工 分 录（一借一贷 · 调账、暂估、计提、冲销用 · 平时一般不用）', MJ_CHK, C_IN,
          '💡 日常的收付款记【资金日记账】、送货单记【送货单登记】、外发记【外发加工登记】、工资记【工资登记】，都会自动生成分录，不用在这里记。'
          '这张表只在要「调账」时用：比如月底暂估还没拿到单子的材料、计提税费、改正以前记错的、冲销。一行一借一贷：填日期、摘要、借方科目编码、贷方科目编码、金额，科目名称自动带出。'
          '金额可以填负数，表示冲回（红字）。动到应收 / 应付 / 其他应收 / 其他应付的，「往来单位」要填，对账单才看得到；动到制造费用、管理 / 销售 / 财务费用的，「费用项目」最好填。'
          '借方或贷方是 400101 / 400102 / 400103 / 4101（生产成本、制造费用）的，会进【成本分摊表】按交货双数分到各款。'
          '400104、540101～540104 是系统月末自动结转的，不能在这里记。记错的行清空内容，不要删行插行。')

    # 第 3 行：汇总
    _kpi(ws, 'A3', '笔数', 'B3', f'=COUNT({mjr(MJ_SEQ)})', CNT)
    _kpi(ws, 'C3', '✗ 待改', 'D3', f'=COUNTIF({mjr(MJ_CHK)},"✗*")', CNT, F_CNT)
    _kpi(ws, 'E3', '⚠ 提醒', 'F3', f'=COUNTIF({mjr(MJ_CHK)},"⚠*")', CNT, F_CNT)
    _kpi(ws, 'G3', '能记账的\n金额合计', 'H3', f'=SUMIFS({mjr(MJ_AMT)},{mjr(MJ_OK)},1)', MONEY)
    fund = '+'.join(f'(LEFT({mjr(x)},4)="{p}")' for x in (MJ_DR, MJ_CR) for p in ('1001', '1002', '1012'))
    _kpi(ws, 'I3', '动了资金科目', 'J3', f'=SUMPRODUCT(({mjr(MJ_OK)}=1)*(({fund})>0))', CNT, F_CNT)
    put(ws, 'K3', '✗ 的行不进账；金额填负数＝冲回（红字）', F_NOTE, KPI_FILL, align=ALW)
    ws.merge_cells('K3:L3')
    ws.row_dimensions[3].height = 30
    ws.conditional_formatting.add('D3', FormulaRule(formula=['N(D3)>0'], fill=FILL_WARN))
    for c in ('F3', 'J3'):
        ws.conditional_formatting.add(c, FormulaRule(formula=[f'N({c})>0'], fill=FILL_YEL))

    heads = [(MJ_SEQ, '序号'), (MJ_DATE, '日期'), (MJ_MEMO, '摘要'), (MJ_DRIN, '借方\n科目编码'), (MJ_DRN, '借方科目'),
             (MJ_CRIN, '贷方\n科目编码'), (MJ_CRN, '贷方科目'), (MJ_AMTIN, '金额\n（负数＝冲回）'), (MJ_UNIT, '往来单位'),
             (MJ_FEE, '费用项目'), (MJ_NOTE, '备注'), (MJ_CHK, '校验')]
    auto_cols = (MJ_SEQ, MJ_DRN, MJ_CRN)
    for col, t in heads:
        c = HDR_CHK if col == MJ_CHK else (HDR_AUTO if col in auto_cols else HDR_IN)
        put(ws, f'{col}{MJ_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[MJ_HDR].height = 34
    hidden = [(MJ_CM, 'CM 月份'), (MJ_DR, 'DR 借方'), (MJ_CR, 'CR 贷方'), (MJ_COMP, 'COMP 成本组件'), (MJ_CAMT, 'CAMT 成本额'),
              (MJ_CKEY, 'CKEY 成本键'), (MJ_OK, 'OK 能记账'), (MJ_AMT, 'AMT 金额'), (X_DRI, '借方科目行'), (X_CRI, '贷方科目行')]
    for col, t in hidden:
        ws[f'{col}{MJ_HDR}'] = t
        ws[f'{col}{MJ_HDR}'].font = F_HELP

    yr, op = P['YEAR'], P['OPEN']
    for r in range(MJ_R0, MJ_R1 + 1):
        A, B, D, F, H, I, J = (f'{c}{r}' for c in (MJ_SEQ, MJ_DATE, MJ_DRIN, MJ_CRIN, MJ_AMTIN, MJ_UNIT, MJ_FEE))
        Lc = f'{MJ_CHK}{r}'
        O, Pp, Q, T, U, DRI, CRI = (f'{c}{r}' for c in (MJ_DR, MJ_CR, MJ_COMP, MJ_OK, MJ_AMT, X_DRI, X_CRI))
        f = {}
        f[MJ_SEQ] = (f'=IF(COUNTA({MJ_DATE}{r}:{MJ_DRIN}{r},{MJ_CRIN}{r},{MJ_AMTIN}{r}:{MJ_NOTE}{r})=0,"",ROW()-{MJ_HDR})')
        f[MJ_CM] = '=' + ym_month(B, yr)
        f[MJ_DR] = f'=TRIM({D}&"")'
        f[MJ_CR] = f'=TRIM({F}&"")'
        f[MJ_AMT] = f'=ROUND(N({H}),2)'
        f[X_DRI] = f'=IF({O}="",0,IFERROR(MATCH({O},{COA_CODES},0),0))'
        f[X_CRI] = f'=IF({Pp}="",0,IFERROR(MATCH({Pp},{COA_CODES},0),0))'
        f[MJ_DRN] = f'=IF({O}="","",IF({DRI}=0,"？",INDEX({COA_NAMES_R},{DRI})&""))'
        f[MJ_CRN] = f'=IF({Pp}="","",IF({CRI}=0,"？",INDEX({COA_NAMES_R},{CRI})&""))'
        cd, cc = comp_of(O), comp_of(Pp)
        f[MJ_COMP] = f'=IF({cd}<>"",{cd},{cc})'
        f[MJ_CAMT] = f'=IF({T}<>1,0,IF({cd}<>"",{U},IF({cc}<>"",-{U},0)))'
        f[MJ_CKEY] = f'=IF(AND({Q}<>"",{T}=1),"公||"&{Q},"")'
        f[MJ_OK] = f'=IF(AND({Lc}<>"",LEFT({Lc},1)<>"✗"),1,0)'

        def any_of(codes):
            return 'OR(' + ','.join(f'{x}="{c}"' for x in (O, Pp) for c in codes) + ')'
        arap = ('2202', '1122', '1221', '2241')
        fee_ac = (MOH, '5601', '5602', '5603')
        fund = 'OR(' + ','.join(f'LEFT({x},4)="{p}"' for x in (O, Pp) for p in ('1001', '1002', '1012')) + ')'
        fei = f'IFERROR(MATCH(TRIM({J}&""),{FE_NAMES},0),0)'
        fcode = f'TRIM(INDEX({FE_CODES},{fei})&"")'
        chk = [
            (f'NOT(ISNUMBER({B}))', '"✗ 日期要填成日期（如 2026/9/30）"'),
            (f'YEAR({B})<>{yr}', '"✗ 不在本会计年度（年份看【基础资料】）"'),
            (f'{B}<{op}', '"✗ 早于建账日期（建账前的数放【会计科目表】年初余额）"'),
            (f'{O}=""', '"✗ 没填借方科目"'),
            (f'{Pp}=""', '"✗ 没填贷方科目"'),
            (f'{DRI}=0', f'"✗ 借方科目 "&{O}&" 不在【会计科目表】：从下拉选，或先去加科目"'),
            (f'{CRI}=0', f'"✗ 贷方科目 "&{Pp}&" 不在【会计科目表】：从下拉选，或先去加科目"'),
            (f'N(INDEX({COA_LEAFS},{DRI}))<>1', f'"✗ 借方 "&{O}&" 是上级科目：要用末级科目（看【会计科目表】「末级」=1）"'),
            (f'N(INDEX({COA_LEAFS},{CRI}))<>1', f'"✗ 贷方 "&{Pp}&" 是上级科目：要用末级科目（看【会计科目表】「末级」=1）"'),
            (f'{O}={Pp}', '"✗ 借贷科目一样"'),
            (f'AND({H}<>"",NOT(ISNUMBER({H})))', '"✗ 金额要填数字"'),
            (f'{U}=0', '"✗ 没有金额"'),
            (any_of(_SYS_CODES), '"✗ 400104、540101～540104 是系统月末自动结转的，不能手工记"'),
            (f'AND({cd}<>"",{cc}<>"")', '"✗ 借贷两边都是成本科目（400101/400102/400103/4101）：拆成两笔，中间用一个往来科目过渡"'),
            (f'AND({any_of(arap)},TRIM({I}&"")="")', '"⚠ 应付/应收/其他应收/其他应付没填往来单位：对账单里看不到这一笔"'),
            (f'AND({any_of(arap)},COUNTIF({UN_NAMES},TRIM({I}&""))=0)', '"⚠ 往来单位没在【往来单位】登记：对账单里看不到这一笔"'),
            (f'AND({any_of(fee_ac)},TRIM({J}&"")="")', '"⚠ 费用科目没填费用项目：费用汇总里算「未分项目」"'),
            (f'AND(TRIM({J}&"")<>"",{fei}=0)', '"⚠ 费用项目不在【基础资料】④：费用汇总里对不上"'),
            (f'AND(TRIM({J}&"")<>"",{fcode}<>{O},{fcode}<>{Pp})',
             f'"⚠ 费用项目「"&TRIM({J}&"")&"」的科目是 "&{fcode}&"，跟借贷科目对不上：费用汇总里算不到"'),
            (fund, '"⚠ 动了资金科目：资金日记账的余额里没有这一笔，尽量在资金日记账记"'),
        ]
        f[MJ_CHK] = f'=IF({A}="","",' + chain(chk, '"√"') + ')'
        for col, v in f.items():
            ws[f'{col}{r}'] = v

    cols = [MJ_SEQ, MJ_DATE, MJ_MEMO, MJ_DRIN, MJ_DRN, MJ_CRIN, MJ_CRN, MJ_AMTIN, MJ_UNIT, MJ_FEE, MJ_NOTE, MJ_CHK]
    in_cols = [c for c in cols if c not in auto_cols and c != MJ_CHK]
    style_rows(ws, MJ_R0, MJ_R1, cols, auto=list(auto_cols) + [MJ_CHK],
               fmts={MJ_SEQ: INT, MJ_DATE: DATE, MJ_DRIN: '@', MJ_CRIN: '@', MJ_AMTIN: MONEY},
               aligns={MJ_MEMO: AL, MJ_DRN: AL, MJ_CRN: AL, MJ_UNIT: AL, MJ_NOTE: AL, MJ_CHK: AL, MJ_AMTIN: AR},
               fills={c: FILL_IN for c in in_cols})
    for r in range(MJ_R0, MJ_R1 + 1):
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hidden])

    dv_date(ws, f'{MJ_DATE}{MJ_R0}:{MJ_DATE}{MJ_R1}')
    dv_list(ws, f'{MJ_DRIN}{MJ_R0}:{MJ_DRIN}{MJ_R1}', f'={COA_CODES}', '科目编码（【会计科目表】的末级科目）')
    dv_list(ws, f'{MJ_CRIN}{MJ_R0}:{MJ_CRIN}{MJ_R1}', f'={COA_CODES}', '科目编码（【会计科目表】的末级科目）')
    dv_list(ws, f'{MJ_UNIT}{MJ_R0}:{MJ_UNIT}{MJ_R1}', f'={UN_NAMES}', '动到应收 / 应付 / 其他应收 / 其他应付的要填', stop=False)
    dv_list(ws, f'{MJ_FEE}{MJ_R0}:{MJ_FEE}{MJ_R1}', f'={FE_NAMES}', '动到制造费用、管理 / 销售 / 财务费用的最好填')
    _check_fmt(ws, f'{MJ_CHK}{MJ_R0}:{MJ_CHK}{MJ_R1}', f'${MJ_CHK}{MJ_R0}')
    for c in (MJ_DRN, MJ_CRN):
        ws.conditional_formatting.add(f'{c}{MJ_R0}:{c}{MJ_R1}', FormulaRule(formula=[f'${c}{MJ_R0}="？"'], fill=FILL_WARN, font=F_REDB))
    ws.auto_filter.ref = f'{MJ_SEQ}{MJ_HDR}:{MJ_CHK}{MJ_R1}'
    ws.freeze_panes = f'{MJ_MEMO}{MJ_R0}'

    for i, row in enumerate(ctx.get(SH_MJ, [])):
        r = MJ_R0 + i
        assert r <= MJ_R1
        for c, v in row.items():
            if v is not None:
                ws[f'{c}{r}'] = v
    return ws


def build(wb, ctx):
    build_cash(wb, ctx)
    build_mj(wb, ctx)
