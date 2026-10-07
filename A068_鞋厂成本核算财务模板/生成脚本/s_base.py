# -*- coding: utf-8 -*-
"""档案几张：【基础资料】【往来单位】【款式档案】【品名档案】【会计科目表】，和隐藏的【_辅助】。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from common import *
from layout import *


def _dv_whole(ws, sqref, lo, hi, msg):
    dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, errorTitle='数字', error=msg)
    ws.add_data_validation(dv)
    dv.add(sqref)


def _band(ws, c1, c2, text, color='FF595959'):
    section(ws, 3, c1, c2, text, color)


F_WIP = Font(name=YH, sz=10, color='FF9C5700')      # 挂着没转 >0：深黄字（配黄底）


def _fit_h(ws, r, text, per_line, line_pt, pad=4, floor=15):
    """按字数把行加高到放得下（只加高不压低：几块并排共用同一行，谁要的高就按谁）"""
    n = max(1, -(-len(str(text or '')) // per_line))
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, floor, n * line_pt + pad)


# ═══════════════════════════ 基础资料 ═══════════════════════════
def build_base(wb, ctx):
    ws = wb[SH_BASE]
    widths(ws, {'A': 18, 'B': 14, 'C': 44, 'D': 2,
                'E': 5, 'F': 12, 'G': 11, 'H': 9, 'I': 13, 'J': 18, 'K': 13, 'L': 13, 'M': 11, 'N': 2,
                'O': 5, 'P': 12, 'Q': 14, 'R': 10, 'S': 24, 'T': 46, 'U': 2,
                'V': 5, 'W': 13, 'X': 11, 'Y': 9, 'Z': 30, 'AA': 2,
                'AB': 5, 'AC': 11, 'AD': 10, 'AE': 30, 'AF': 2,
                'AG': 5, 'AH': 12, 'AI': 30, 'AJ': 2,
                'AK': 5, 'AL': 22, 'AM': 12, 'AN': 12, 'AO': 9, 'AP': 8, 'AQ': 9, 'AR': 13, 'AS': 11, 'AT': 30, 'AU': 2,
                'BI': 5, 'BJ': 26, 'BK': 8, 'BL': 2, 'BM': 12})
    title(ws, '基 础 资 料（参数 · 资金账户 · 收支类别 · 费用项目 · 部门 · 材料类别 · 固定资产）', 'T', C_ARC,
          '💡 一年开始时设好，平时很少动。淡黄格子是可以改的；灰色是自动算的，或者是预置的（③ 收支类别前 17 行的名称、对方科目、还要选：公式按名字认，不要改）。'
          '几块并排放着，往右拉能看到全部（⑦ 固定资产在 AK 列起）；① 下面是 ⑩ 每月结账日期。'
          '每块只往下面的空行加，不要插行删行。资金账户最多 8 个（银行最多 6 个、支付宝微信最多 4 个、「现金」类型只能 1 个）；改了账户名称，资金日记账里已经记的要跟着改。')
    ws.row_dimensions[2].height = 48
    for c1, c2, t in (('A', 'C', '① 参数'), ('E', 'M', '② 资金账户（期初＝建账日那天的余额）'), ('O', 'T', '③ 收支类别（资金日记账用）　灰色是预置的：名称、对方科目、还要选不要改；新类别往下面空行加'),
                      ('V', 'Z', '④ 费用项目'), ('AB', 'AE', '⑤ 部门（工资登记用）'), ('AG', 'AI', '⑥ 材料类别'),
                      ('AK', 'AT', '⑦ 固定资产（自动按月提折旧）'), ('BI', 'BK', '⑧ 现金流量项目'), ('BM', 'BM', '⑨ 外发工序')):
        _band(ws, c1, c2, t)
    hdr = {'A': '项目', 'B': '值', 'C': '说明',
           AC_SEQ: '序号', AC_NAME: '账户名称', AC_TYPE: '类型', AC_CODE: '科目编码\n（自动）', AC_OPEN: '期初余额', AC_NO: '账号 / 说明',
           AC_NOW: '当前余额\n（按日记账）', AC_REAL: '实际余额\n（手填核对）', AC_DIFF: '差额',
           CT_SEQ: '序号', CT_NAME: '收支类别', CT_OPP: '对方科目', CT_NEED: '还要选', CT_CF: '现金流量项目', CT_NOTE: '什么时候用',
           FE_SEQ: '序号', FE_NAME: '费用项目', FE_CLS: '归类', FE_CODE: '科目\n（自动）', FE_NOTE: '说明',
           DP_SEQ: '序号', DP_NAME: '部门', DP_TYPE: '类型', DP_NOTE: '工资进哪里（自动）',
           MC_SEQ: '序号', MC_NAME: '材料类别', MC_NOTE: '包括',
           FA_SEQ: '序号', FA_NAME: '名称', FA_DATE: '购入日期', FA_COST: '原值', FA_MON: '折旧月数', FA_RES: '残值率',
           FA_USE: '使用部门', FA_DEP0: '建账前\n已提折旧', FA_DEPM: '月折旧\n（自动）', FA_NOTE: '备注',
           CF_SEQ: '序号', CF_NAME: '项目', CF_GRP: '大类', OP_NAME: '工序'}
    for col, t in hdr.items():
        put(ws, f'{col}{B_HDR}', t, F_HDR, fill('FF595959'), align=ACW)
    ws.row_dimensions[B_HDR].height = 32

    # ① 参数
    M12 = '{1,2,3,4,5,6,7,8,9,10,11,12}'     # 空行的月份列可能是 ""：用 COUNTIFS 按月数，不拿月份列直接相乘
    last_m = ('MAX(' + ','.join(f'SUMPRODUCT(MAX((COUNTIFS({a},{M12},{o},1)>0)*{M12}))' for a, o in (
        (jr(J_CM), jr(J_OK)), (dnr(DN_CM), dnr(DN_OK)), (otr(OT_CM), otr(OT_OK)), (odr(OD_DM), odr(OD_OK)),
        (wgr(WG_CM), wgr(WG_OK)), (mjr(MJ_CM), mjr(MJ_OK)))) + ')')
    auto = {
        '=auto': f'=IF(N({P["DEPTO"]})>=1,MIN(12,INT(N({P["DEPTO"]}))),N({P["LASTM"]}))',
        '=lockauto': f'=SUMPRODUCT(MAX(ISNUMBER({CLOSE_DATES})*(ROW({CLOSE_DATES})-{CL_R0 - 1})))',
        '=lastm': f'=MAX(IF(YEAR({P["OPEN"]})={P["YEAR"]},MONTH({P["OPEN"]}),1),MIN(12,{last_m}))',
    }
    # 参数说明按这一轮的口径（layout 里只改了 LOCK 那句；这几句写到表上时换掉）
    note_fix = {
        'LAGD': f'收单日期比送货日期晚这么多天以上的，【{SH_LATE}】列出来',
        'DEPTO': '空着＝自动：提到「最新月份」（最下面一格：有业务的最后一个月），后面还没到的月份先不提；要提到某个月就填几',
        'DEPTOX': '自动：上一格填了就按上一格；空着＝最新月份',
    }
    for k, lbl, val, fmt, note in PARAMS:
        r = PA_ROW[k]
        note = note_fix.get(k, note)
        put(ws, f'{PA_LBL}{r}', lbl, F_TXTB, align=AL)
        if val in auto:
            put(ws, f'{PA_VAL}{r}', auto[val], F_AUTO, FILL_AUTO, fmt=fmt, align=AC)
        else:
            put(ws, f'{PA_VAL}{r}', val, F_IN, FILL_IN, fmt=fmt, align=AC)
        put(ws, f'{PA_NOTE}{r}', note, F_NOTE, align=ALW)
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 15 * (-(-len(note) // 22)) + 6)
    # ⑩ 每月结账日期
    section(ws, CL_HDR, 'A', 'C', '⑩ 每月结账日期（只管送货单 / 外发单晚到）', 'FF595959')
    for m in range(1, 13):
        r = CL_R0 + m - 1
        put(ws, f'A{r}', f'{m}月', F_TXTB, align=AC)
        put(ws, f'{CL_DATE}{r}', None, F_IN, FILL_IN, fmt=DATE, align=AC)
        put(ws, f'C{r}', None, F_NOTE, align=ALW)
    ws[f'C{CL_R0}'] = ('x 月的报表已经报给老板、不想再变了，就在 x 月这一行填结账那天（比如 8 月填 2026/9/5）。'
                       '以后才收到的 x 月送货单 / 外发单，算到收单那个月，x 月的数不再变。填了就不要改、不要清。')
    ws.merge_cells(f'C{CL_R0}:C{CL_R0 + 5}')
    ws[f'C{CL_R0 + 6}'] = '只管供应商单据晚到；工资、交货、收付款、手工分录补录到已结账的月份，照样会改动那个月。'
    ws.merge_cells(f'C{CL_R0 + 6}:C{CL_R1}')
    dv_date(ws, f'{CL_DATE}{CL_R0}:{CL_DATE}{CL_R1}')
    dv_list(ws, P['BASIS'].split('!')[1].replace('$', ''), '"系数,标准成本"', '系数 或 标准成本')
    dv_list(ws, P['PERIOD'].split('!')[1].replace('$', ''), '"按收入,不摊"', '按收入 或 不摊')
    dv_list(ws, P['CUST'].split('!')[1].replace('$', ''), f'={UN_NAMES}', '【往来单位】里的客户', stop=False)
    _dv_whole(ws, P['YEAR'].split('!')[1].replace('$', ''), 2000, 2100, '填年份，比如 2026')
    _dv_whole(ws, P['LAGD'].split('!')[1].replace('$', ''), 1, 999, '填天数')
    _dv_whole(ws, P['DEPTO'].split('!')[1].replace('$', ''), 1, 12, '填 1～12，或者空着')
    dv_date(ws, P['OPEN'].split('!')[1].replace('$', ''))

    # ② 资金账户
    accs = ctx.get('accounts', [])
    for i, r in enumerate(range(AC_R0, AC_R1 + 1)):
        ws[f'{AC_SEQ}{r}'] = f'=IF({AC_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(accs):
            nm, tp, op, note = accs[i]
            ws[f'{AC_NAME}{r}'], ws[f'{AC_TYPE}{r}'], ws[f'{AC_OPEN}{r}'], ws[f'{AC_NO}{r}'] = nm, tp, op, note
        g = f'{AC_TYPE}{r}'
        ws[f'{AC_CODE}{r}'] = (f'=IF({AC_NAME}{r}="","",IF({g}="现金","1001",'
                               f'IF({g}="银行",IF(COUNTIF(${AC_TYPE}${AC_R0}:{g},"银行")>6,"","1002"&TEXT(COUNTIF(${AC_TYPE}${AC_R0}:{g},"银行"),"00")),'
                               f'IF({g}="支付宝微信",IF(COUNTIF(${AC_TYPE}${AC_R0}:{g},"支付宝微信")>4,"","1012"&TEXT(COUNTIF(${AC_TYPE}${AC_R0}:{g},"支付宝微信"),"00")),""))))')
        nm = f'{AC_NAME}{r}'
        ws[f'{AC_NOW}{r}'] = (f'=IF({nm}="","",ROUND(N({AC_OPEN}{r})+SUMIFS({jr(J_NET)},{jr(J_ACC)},{nm})'
                              f'-SUMIFS({jr(J_NET)},{jr(J_TOX)},{nm}),2))')
        ws[f'{AC_DIFF}{r}'] = f'=IF(OR({nm}="",{AC_REAL}{r}=""),"",ROUND(N({AC_REAL}{r})-N({AC_NOW}{r}),2))'
    style_rows(ws, AC_R0, AC_R1, [AC_SEQ, AC_NAME, AC_TYPE, AC_CODE, AC_OPEN, AC_NO, AC_NOW, AC_REAL, AC_DIFF],
               auto=[AC_SEQ, AC_CODE, AC_NOW, AC_DIFF],
               fmts={AC_OPEN: MONEY, AC_NOW: MONEY, AC_REAL: MONEY, AC_DIFF: MONEY, AC_CODE: '@'},
               aligns={AC_NO: AL, AC_OPEN: AR, AC_NOW: AR, AC_REAL: AR, AC_DIFF: AR},
               fills={c: FILL_IN for c in (AC_NAME, AC_TYPE, AC_OPEN, AC_NO, AC_REAL)})
    dv_list(ws, f'{AC_TYPE}{AC_R0}:{AC_TYPE}{AC_R1}', '"' + ','.join(AC_TYPES) + '"', '银行 / 现金 / 支付宝微信')
    ws.conditional_formatting.add(f'{AC_DIFF}{AC_R0}:{AC_DIFF}{AC_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER({AC_DIFF}{AC_R0}),ABS({AC_DIFF}{AC_R0})>=0.01)'], fill=FILL_WARN, font=F_RED))
    put(ws, f'{AC_NAME}{AC_R1 + 2}', '账户合计', F_TXTB, FILL_SUB, align=AC)
    put(ws, f'{AC_OPEN}{AC_R1 + 2}', f'=SUM({AC_OPEN}{AC_R0}:{AC_OPEN}{AC_R1})', F_AUTOB, FILL_SUB, MONEY, AR)
    put(ws, f'{AC_NOW}{AC_R1 + 2}', f'=SUM({AC_NOW}{AC_R0}:{AC_NOW}{AC_R1})', F_AUTOB, FILL_SUB, MONEY, AR)
    put(ws, f'{AC_NAME}{AC_R1 + 3}', '「当前余额」＝期初＋资金日记账里这个账户的收支（含内部转账）。「实际余额」照银行 App / 支付宝填，差额不是 0 就是漏记或记错了。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'{AC_NAME}{AC_R1 + 3}:{AC_DIFF}{AC_R1 + 4}')

    # ③ 收支类别：前 17 行预置（名称、对方科目、还要选是公式按名字认的 → 灰色，不要改）；第 CT_R0+17 行起用户自己加（淡黄）
    for i, r in enumerate(range(CT_R0, CT_R1 + 1)):
        ws[f'{CT_SEQ}{r}'] = f'=IF({CT_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(CATS):
            nm, opp, need, cf, note = CATS[i]
            ws[f'{CT_NAME}{r}'] = nm
            ws[f'{CT_OPP}{r}'] = {'@费用': '（按费用项目）', '@账户': '（对方账户）'}.get(opp, opp)
            ws[f'{CT_NEED}{r}'] = need
            ws[f'{CT_CF}{r}'] = {'@费用': '（按费用项目）'}.get(cf, cf) or None
            ws[f'{CT_NOTE}{r}'] = note
    style_rows(ws, CT_R0, CT_R1, [CT_SEQ, CT_NAME, CT_OPP, CT_NEED, CT_CF, CT_NOTE], auto=[CT_SEQ],
               fmts={CT_OPP: '@'}, aligns={CT_NOTE: ALW, CT_CF: ALW},
               fills={c: FILL_IN for c in (CT_NAME, CT_OPP, CT_NEED, CT_CF, CT_NOTE)})
    ct_user0 = CT_R0 + len(CATS)                 # 用户自己加的第一行（第 22 行）
    for r in range(CT_R0, ct_user0):
        for c in (CT_NAME, CT_OPP, CT_NEED):
            ws[f'{c}{r}'].fill = FILL_AUTO
            ws[f'{c}{r}'].font = F_TXT
    put(ws, f'{CT_SEQ}{CT_R1 + 1}', '灰色是预置的 17 个类别：名称、对方科目、还要选不要改（公式按名字认，改了账户余额、自动认类别都会出错）；'
                                     '现金流量项目、说明可以改。要新的类别往下面空行加（第 %d 行起），名称别跟预置的重复。' % ct_user0,
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'{CT_SEQ}{CT_R1 + 1}:{CT_NOTE}{CT_R1 + 2}')
    for r in range(CT_R0, ct_user0):
        _fit_h(ws, r, ws[f'{CT_NOTE}{r}'].value, 22, 14, floor=30)
    dv_list(ws, f'{CT_NEED}{ct_user0}:{CT_NEED}{CT_R1}', '"单位,费用项目,对方账户,无"', '选了这个类别以后还必须选什么')
    dv_list(ws, f'{CT_CF}{CT_R0}:{CT_CF}{CT_R1}', f'={CF_NAMES}', '现金流量表里算哪一项', stop=False)

    # ④ 费用项目
    for i, r in enumerate(range(FE_R0, FE_R1 + 1)):
        ws[f'{FE_SEQ}{r}'] = f'=IF({FE_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(FEES):
            ws[f'{FE_NAME}{r}'], ws[f'{FE_CLS}{r}'], ws[f'{FE_NOTE}{r}'] = FEES[i][0], FEES[i][1], FEES[i][2] or None
        x = f'{FE_CLS}{r}'
        f = '""'
        for k, (code, _cf) in reversed(list(FE_CLASSES.items())):
            f = f'IF({x}="{k}","{code}",{f})'
        ws[f'{FE_CODE}{r}'] = f'=IF({FE_NAME}{r}="","",{f})'
    style_rows(ws, FE_R0, FE_R1, [FE_SEQ, FE_NAME, FE_CLS, FE_CODE, FE_NOTE], auto=[FE_SEQ, FE_CODE],
               fmts={FE_CODE: '@'}, aligns={FE_NOTE: ALW}, fills={c: FILL_IN for c in (FE_NAME, FE_CLS, FE_NOTE)})
    for i in range(len(FEES)):
        if len(FEES[i][2] or '') > 14:          # 说明列 30 宽，一行放 14 个字左右；放不下的才加高
            _fit_h(ws, FE_R0 + i, FEES[i][2], 14, 14)
    dv_list(ws, f'{FE_CLS}{FE_R0}:{FE_CLS}{FE_R1}', '"' + ','.join(FE_CLASSES) + '"',
            '制造费用＝车间里花的（会分到款式成本里）；管理/销售/财务费用＝不分到款式，直接进利润表')

    # ⑤ 部门
    for i, r in enumerate(range(DP_R0, DP_R1 + 1)):
        ws[f'{DP_SEQ}{r}'] = f'=IF({DP_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(DEPTS):
            ws[f'{DP_NAME}{r}'], ws[f'{DP_TYPE}{r}'] = DEPTS[i]
        x = f'{DP_TYPE}{r}'
        ws[f'{DP_NOTE}{r}'] = (f'=IF({DP_NAME}{r}="","",IF({x}="直接生产","生产成本-直接人工（计件可按款式记）",'
                               f'IF({x}="车间辅助","制造费用（分到各款式）",IF({x}="管理","管理费用",IF({x}="销售","销售费用","✗ 选类型")))))')
    style_rows(ws, DP_R0, DP_R1, [DP_SEQ, DP_NAME, DP_TYPE, DP_NOTE], auto=[DP_SEQ, DP_NOTE], aligns={DP_NOTE: AL},
               fills={c: FILL_IN for c in (DP_NAME, DP_TYPE)})
    dv_list(ws, f'{DP_TYPE}{DP_R0}:{DP_TYPE}{DP_R1}', '"' + ','.join(DP_TYPES) + '"', '直接生产 / 车间辅助 / 管理 / 销售')

    # ⑥ 材料类别
    for i, r in enumerate(range(MC_R0, MC_R1 + 1)):
        ws[f'{MC_SEQ}{r}'] = f'=IF({MC_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(MATCATS):
            ws[f'{MC_NAME}{r}'], ws[f'{MC_NOTE}{r}'] = MATCATS[i][0], MATCATS[i][1] or None
    style_rows(ws, MC_R0, MC_R1, [MC_SEQ, MC_NAME, MC_NOTE], auto=[MC_SEQ], aligns={MC_NOTE: ALW},
               fills={c: FILL_IN for c in (MC_NAME, MC_NOTE)})

    # ⑦ 固定资产
    fas = ctx.get('fa', [])
    yr, op = P['YEAR'], P['OPEN']
    t_open = f'(YEAR({op})*12+MONTH({op}))'
    for i, r in enumerate(range(FA_R0, FA_R1 + 1)):
        ws[f'{FA_SEQ}{r}'] = f'=IF({FA_NAME}{r}="","",ROW()-{B_HDR})'
        if i < len(fas):
            for c, v in zip((FA_NAME, FA_DATE, FA_COST, FA_MON, FA_RES, FA_USE, FA_DEP0, FA_NOTE), fas[i]):
                ws[f'{c}{r}'] = v
        cost, mon, res, dep0, dm = (f'${c}{r}' for c in (FA_COST, FA_MON, FA_RES, FA_DEP0, FA_DEPM))
        ws[f'{FA_DEPM}{r}'] = f'=IF(OR(${FA_NAME}{r}="",N({cost})=0,N({mon})<=0),"",ROUND(N({cost})*(1-N({res}))/N({mon}),2))'
        start = f'(YEAR(${FA_DATE}{r})*12+MONTH(${FA_DATE}{r})+1)'
        for m in range(1, 13):
            t = f'({yr}*12+{m})'
            ws[f'{fa_mcol(m)}{r}'] = (f'=IF(OR(${FA_NAME}{r}="",NOT(ISNUMBER(${FA_DATE}{r})),N({dm})=0,{m}>N({P["DEPTOX"]})),0,'
                                      f'IF(OR({t}<{t_open},{t}<{start}),0,ROUND(MAX(0,MIN(N({dm}),'
                                      f'N({cost})*(1-N({res}))-N({dep0})-N({dm})*({t}-MAX({start},{t_open})))),2)))')
            ws[f'{fa_mcol(m)}{r}'].font = F_HELP
    for m in range(1, 13):
        ws[f'{fa_mcol(m)}{B_HDR}'] = f'{m}月折旧'
        ws[f'{fa_mcol(m)}{B_HDR}'].font = F_HELP
    hide(ws, *[fa_mcol(m) for m in range(1, 13)], 'AU')
    style_rows(ws, FA_R0, FA_R1, [FA_SEQ, FA_NAME, FA_DATE, FA_COST, FA_MON, FA_RES, FA_USE, FA_DEP0, FA_DEPM, FA_NOTE],
               auto=[FA_SEQ, FA_DEPM], fmts={FA_DATE: DATE, FA_COST: MONEY, FA_MON: INT, FA_RES: '0%', FA_DEP0: MONEY, FA_DEPM: MONEY},
               aligns={FA_NAME: AL, FA_NOTE: AL, FA_COST: AR, FA_DEP0: AR, FA_DEPM: AR},
               fills={c: FILL_IN for c in (FA_NAME, FA_DATE, FA_COST, FA_MON, FA_RES, FA_USE, FA_DEP0, FA_NOTE)})
    dv_list(ws, f'{FA_USE}{FA_R0}:{FA_USE}{FA_R1}', '"' + ','.join(FA_USES) + '"', '车间＝折旧进制造费用（分到款式）；管理/销售＝进期间费用')
    dv_date(ws, f'{FA_DATE}{FA_R0}:{FA_DATE}{FA_R1}')
    put(ws, f'{FA_NAME}{FA_R1 + 2}', '规则：购入的下个月开始提，每月＝原值×(1−残值率)÷折旧月数，提满为止；建账日期以前的月份不提（已经提的填「建账前已提折旧」）。'
                                     '建账以前买的：原值、已提折旧自动进期初；建账以后买的：资金日记账记「购置设备」，再在这里登记一行。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'{FA_NAME}{FA_R1 + 2}:{FA_NOTE}{FA_R1 + 4}')

    # ⑧ 现金流量项目 ⑨ 外发工序
    for i, (nm, grp) in enumerate(CF_ITEMS):
        r = CF_R0 + i
        ws[f'{CF_SEQ}{r}'], ws[f'{CF_NAME}{r}'], ws[f'{CF_GRP}{r}'] = i + 1, nm, grp
    style_rows(ws, CF_R0, CF_R1, [CF_SEQ, CF_NAME, CF_GRP], auto=[CF_SEQ, CF_NAME, CF_GRP], aligns={CF_NAME: AL})
    for i, r in enumerate(range(OP_R0, OP_R1 + 1)):
        if i < len(OPS):
            ws[f'{OP_NAME}{r}'] = OPS[i]
    style_rows(ws, OP_R0, OP_R1, [OP_NAME], fills={OP_NAME: FILL_IN})
    ws.freeze_panes = 'A5'


# ═══════════════════════════ 往来单位 ═══════════════════════════
def build_unit(wb, ctx):
    ws = wb[SH_UNIT]
    widths(ws, {'A': 5, 'B': 20, 'C': 11, 'D': 9, 'E': 13, 'F': 26, 'G': 8, 'H': 12, 'I': 12, 'J': 12, 'K': 12, 'L': 22, 'M': 22})
    title(ws, '往 来 单 位（供应商 · 外发加工厂 · 客户 · 老板个人）', 'M', C_ARC,
          '💡 一家一行，名称要跟送货单、资金日记账里写的一模一样（下拉就是从这里来的）。类型决定资金日记账怎么自动认：'
          '材料供应商付钱＝付材料款，外发加工厂付钱＝付加工费，客户来钱＝收货款。期初余额填建账日那天的数（我们欠他的填「期初应付」，他欠我们的填「期初应收」）。'
          'M 列「现在余额」按全部已记的账算＝（应付＋其他应付）−（应收＋其他应收）：正数＝我们还欠他（供应商、加工厂的货款，老板借给厂里的钱），'
          '负数＝我们付多了 / 他欠我们；类型是「客户」的反过来：正数＝他还欠我们，负数＝他多付了（预收）。')
    put(ws, 'A3', '家数', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'B3', f'=SUMPRODUCT(--(TRIM({UN_NAMES}&"")<>""))', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'C3', '期初应付合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'D3', f'=SUM({unr(UN_AP0)})', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    ws.merge_cells('D3:E3')
    put(ws, 'F3', '期初应收合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'G3', f'=SUM({unr(UN_AR0)})', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    ws.merge_cells('G3:H3')
    heads = [(UN_SEQ, '序号'), (UN_NAME, '名称'), (UN_TYPE, '类型'), (UN_CONT, '联系人'), (UN_TEL, '电话'), (UN_MAIN, '主营 / 说明'),
             (UN_TERM, '结算'), (UN_AP0, '期初应付\n（欠他的）'), (UN_AR0, '期初应收\n（他欠的）'), (UN_OR0, '期初\n其他应收'),
             (UN_OP0, '期初\n其他应付'), (UN_NOTE, '备注'), ('M', '现在余额（自动）\n正数＝我们欠他\n客户：正数＝他欠我们')]
    header(ws, UN_HDR, heads, 'FF595959', height=46)
    hidden = [(UN_BAP, '应付余额@报表月'), (UN_BAR, '应收余额@报表月'), (UN_BOR, '其他应收@报表月'), (UN_BOP, '其他应付@报表月'),
              (UN_APK, '供应商序号键'), (UN_ARK, '客户序号键')]
    for c, t in hidden:
        ws[f'{c}{UN_HDR}'] = t
        ws[f'{c}{UN_HDR}'].font = F_HELP
    rows = ctx.get(SH_UNIT, [])
    m = BS_MX               # 资产负债表实际用的月份（C3 空着＝最新月份）
    for i, r in enumerate(range(UN_R0, UN_R1 + 1)):
        if i < len(rows):
            for c, v in rows[i].items():
                if v is not None:
                    ws[f'{c}{r}'] = v
        nm, tp = f'{UN_NAME}{r}', f'{UN_TYPE}{r}'
        ws[f'{UN_SEQ}{r}'] = f'=IF({nm}="","",ROW()-{UN_HDR})'

        def bal(code, col0, sign, mcrit=None):
            d = je_sum(f'"{code}"', 'D', mcrit, nm)
            c_ = je_sum(f'"{code}"', 'C', mcrit, nm)
            return f'N({col0}{r})+{d}-{c_}' if sign > 0 else f'N({col0}{r})+{c_}-{d}'
        # 现在余额（全部已记的）：(应付 2202＋其他应付 2241) − (应收 1122＋其他应收 1221)，我们欠他为正；客户取反（他欠我们为正）
        #   ＝ 期初(H＋K−I−J) ＋ 四个科目贷方合计 − 四个科目借方合计
        cr4 = '+'.join(je_sum(f'"{c}"', 'C', None, nm) for c in ('2202', '2241', '1122', '1221'))
        dr4 = '+'.join(je_sum(f'"{c}"', 'D', None, nm) for c in ('2202', '2241', '1122', '1221'))
        ws[f'M{r}'] = (f'=IF({nm}="","",ROUND(IF({tp}="客户",-1,1)*(N({UN_AP0}{r})+N({UN_OP0}{r})-N({UN_AR0}{r})-N({UN_OR0}{r})'
                       f'+{cr4}-({dr4})),2))')
        mc = f'"<="&{m}'
        ws[f'{UN_BAP}{r}'] = f'=IF({nm}="",0,ROUND({bal("2202", UN_AP0, -1, mc)},2))'
        ws[f'{UN_BAR}{r}'] = f'=IF({nm}="",0,ROUND({bal("1122", UN_AR0, 1, mc)},2))'
        ws[f'{UN_BOR}{r}'] = f'=IF({nm}="",0,ROUND({bal("1221", UN_OR0, 1, mc)},2))'
        ws[f'{UN_BOP}{r}'] = f'=IF({nm}="",0,ROUND({bal("2241", UN_OP0, -1, mc)},2))'
        ws[f'{UN_APK}{r}'] = f'=IF(AND({nm}<>"",OR({tp}="材料供应商",{tp}="外发加工厂")),ROW()-{UN_HDR},"")'
        ws[f'{UN_ARK}{r}'] = f'=IF(AND({nm}<>"",{tp}="客户"),ROW()-{UN_HDR},"")'
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    cols = [UN_SEQ, UN_NAME, UN_TYPE, UN_CONT, UN_TEL, UN_MAIN, UN_TERM, UN_AP0, UN_AR0, UN_OR0, UN_OP0, UN_NOTE, 'M']
    style_rows(ws, UN_R0, UN_R1, cols, auto=[UN_SEQ, 'M'],
               fmts={UN_AP0: MONEY, UN_AR0: MONEY, UN_OR0: MONEY, UN_OP0: MONEY, 'M': MONEY, UN_TEL: '@'},
               aligns={UN_NAME: AL, UN_MAIN: AL, UN_NOTE: AL, UN_AP0: AR, UN_AR0: AR, UN_OR0: AR, UN_OP0: AR, 'M': AR},
               fills={c: FILL_IN for c in cols if c not in (UN_SEQ, 'M')})
    for r in range(UN_R0, UN_R1 + 1):
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hidden])
    dv_list(ws, f'{UN_TYPE}{UN_R0}:{UN_TYPE}{UN_R1}', '"' + ','.join(UN_TYPES) + '"', '材料供应商 / 外发加工厂 / 客户 / 股东个人 / 其他')
    ws.conditional_formatting.add(f'{UN_NAME}{UN_R0}:{UN_NAME}{UN_R1}', FormulaRule(
        formula=[f'AND({UN_NAME}{UN_R0}<>"",COUNTIF(${UN_NAME}${UN_R0}:${UN_NAME}${UN_R1},{UN_NAME}{UN_R0})>1)'], fill=FILL_WARN, font=F_RED))
    ws.auto_filter.ref = f'A{UN_HDR}:M{UN_R1}'
    ws.freeze_panes = f'C{UN_R0}'


# ═══════════════════════════ 款式档案 ═══════════════════════════
def build_style(wb, ctx):
    ws = wb[SH_STY]
    widths(ws, {'A': 5, 'B': 12, 'C': 16, 'D': 8, 'E': 10, 'F': 8, 'G': 8, 'H': 10, 'I': 10, 'J': 7, 'K': 18,
                'L': 9, 'M': 9, 'N': 9, 'O': 12, 'P': 12, 'Q': 10, 'R': 8, ST_WIP: 12, 'X': 5, 'Y': 12, 'Z': 12, 'AA': 9})
    title(ws, '款 式 档 案（款式编码 · 结算单价 · 分摊系数）', ST_WIP, C_ARC,
          '💡 订单上出现的每个款式编码都要在这里登记一行（右边会自动列出「订单里有、这里还没有」的，照着抄过来就行）。'
          '结算单价＝工厂交给电商部（客户）一双算多少钱，交货金额就按它算；某张订单价钱不一样，在【订单明细】「结算单价」列改那几行。'
          '材料系数 / 工费系数：公共成本（没写用在哪个款的材料、房租水电、计时工资……）按「交货双数×系数」分到各款，默认都是 1（＝按双数平均）；'
          '比如靴子用料是单鞋的 2 倍，材料系数填 2。标准成本可以不填；填了，在【基础资料】把分摊依据改成「标准成本」就按它分，【成本分摊表】还会拿它跟实际比（看出有没有送货单没到）。'
          'S 列「挂着没转」＝直接记到这个款（用在哪＝这个款）、还没转成成本的钱：按交货进度转，还有没交的订单就留一部分等后面那几批（比如开模费记到 HK15，HK15 交货以前就挂在这里）；订单取消、不做了的，把【订单明细】订单数量改成实交数。'
          '款式不再做了，把「状态」改成「停产」：它最后一次交货以后再记到这个款的专用成本（交完货才补来的送货单、外发单等），就按公共分摊，不会一直挂着。')
    ws.row_dimensions[2].height = max(ws.row_dimensions[2].height or 0, 80)      # title() 按字数算过；不够 80 就给 80
    put(ws, 'A3', '款式数', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('A3:B3')
    put(ws, 'C3', f'=SUMPRODUCT(--(TRIM({ST_CODES}&"")<>""))', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'D3', '没登记的款式', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('D3:E3')
    put(ws, 'F3', f'=COUNT({odr(OD_SNEW)})', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'G3', '本年交货双数', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('G3:H3')
    put(ws, 'I3', f'=SUM({str_(ST_DQ)})', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'J3', '本年收入', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('J3:K3')
    put(ws, 'L3', f'=SUM({str_(ST_REV)})', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    ws.merge_cells('L3:M3')
    put(ws, 'N3', '本年成本', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'O3', f'=SUM({str_(ST_COST)})', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    ws.merge_cells('O3:P3')
    put(ws, 'Q3', '挂着没转', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'R3', f'=SUM({str_(ST_WIP)})', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    ws.merge_cells(f'R3:{ST_WIP}3')
    heads = [(ST_SEQ, '序号'), (ST_CODE, '款式编码'), (ST_NAME, '名称 / 说明'), (ST_KIND, '类型'), (ST_PRICE, '结算单价\n（元/双）'),
             (ST_KM, '材料\n系数'), (ST_KL, '工费\n系数'), (ST_SDM, '标准材料\n成本/双'), (ST_SDL, '标准工费\n成本/双'), (ST_STAT, '状态'),
             (ST_NOTE, '备注'), (ST_OQ, '本年订单\n双数'), (ST_DQ, '已交\n双数'), (ST_UQ, '未交\n双数'), (ST_REV, '本年\n交货收入'),
             (ST_COST, '本年成本'), (ST_UNIT, '单双成本'), (ST_GM, '毛利率'), (ST_WIP, '挂着没转\n的直接成本')]
    for col, t in heads:
        c = 'FF595959' if CI(col) <= CI(ST_NOTE) else 'FF548235'
        put(ws, f'{col}{ST_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[ST_HDR].height = 34
    for c, t in ((ST_WM, '材料权重'), (ST_WL, '工费权重'), (ST_KEY, '编码规范写法'), (ST_REL, '停产后转公共的起始月')):
        ws[f'{c}{ST_HDR}'] = t
        ws[f'{c}{ST_HDR}'].font = F_HELP
    rows = ctx.get(SH_STY, [])
    basis = P['BASIS']
    avg_m = f'IFERROR(AVERAGEIF(${ST_SDM}${ST_R0}:${ST_SDM}${ST_R1},">0"),1)'
    avg_l = f'IFERROR(AVERAGEIF(${ST_SDL}${ST_R0}:${ST_SDL}${ST_R1},">0"),1)'
    for i, r in enumerate(range(ST_R0, ST_R1 + 1)):
        if i < len(rows):
            for c, v in rows[i].items():
                if v is not None:
                    ws[f'{c}{r}'] = v
        code = f'{ST_CODE}{r}'
        k = r - ST_R0 + 1
        ws[f'{ST_SEQ}{r}'] = f'=IF({code}="","",ROW()-{ST_HDR})'
        ws[f'{ST_KEY}{r}'] = f'=IF(TRIM({code}&"")="","",TRIM({code}&""))'
        km, kl = f'IF({ST_KM}{r}="",1,N({ST_KM}{r}))', f'IF({ST_KL}{r}="",1,N({ST_KL}{r}))'
        ws[f'{ST_WM}{r}'] = f'=IF({ST_KEY}{r}="",0,IF({basis}="标准成本",IF(N({ST_SDM}{r})>0,{ST_SDM}{r},{km}*{avg_m}),{km}))'
        ws[f'{ST_WL}{r}'] = f'=IF({ST_KEY}{r}="",0,IF({basis}="标准成本",IF(N({ST_SDL}{r})>0,{ST_SDL}{r},{kl}*{avg_l}),{kl}))'
        e = f'{ST_KEY}{r}=""'
        ws[f'{ST_OQ}{r}'] = f'=IF({e},"",SUMIFS({odr(OD_OQ)},{odr(OD_SROW)},{k}))'
        smq = q(SH_SM)
        ws[f'{ST_DQ}{r}'] = f"=IF({e},\"\",SUM({smq}!{sm_blk('双数', r)}))"     # ＝_款式月 双数（订单明细 CQ，不含退货、返修）
        ws[f'{ST_UQ}{r}'] = f'=IF({e},"",MAX(0,N({ST_OQ}{r})-N({ST_DQ}{r})))'
        ws[f'{ST_REV}{r}'] = f"=IF({e},\"\",SUM({q(SH_SM)}!{sm_blk('收入', r)}))"
        ws[f'{ST_COST}{r}'] = f"=IF({e},\"\",SUM({q(SH_SM)}!{sm_blk('成本', r)}))"
        ws[f'{ST_UNIT}{r}'] = f"=IF({e},\"\",IFERROR(ROUND(N({ST_COST}{r})/SUM({q(SH_SM)}!{sm_blk('双数', r)}),2),\"\"))"
        ws[f'{ST_GM}{r}'] = f'=IF({e},"",IF(N({ST_REV}{r})=0,"",ROUND((N({ST_REV}{r})-N({ST_COST}{r}))/N({ST_REV}{r}),4)))'
        # S 挂着没转＝全年直接记到这个款的成本（各来源 CKEY＝"款|编码|*"）− _款式月 转材料～转制造 1～12 月
        key = '"款|"&' + esc(f'{ST_KEY}{r}') + '&"|*"'
        dsum = '+'.join(f'SUMIFS({rng(sh, ca, r0_, r1_)},{rng(sh, ck, r0_, r1_)},{key})' for sh, ck, cm, ca, r0_, r1_, _c in COST_SRC)
        tr = (f"SUM({smq}!{sm_col('转' + COMPS[0], 1)}{r}:{sm_col('转' + COMPS[-1], 12)}{r})"
              f"+SUM({smq}!{sm_col('停' + COMPS[0], 1)}{r}:{sm_col('停' + COMPS[-1], 12)}{r})")
        ws[f'{ST_WIP}{r}'] = f'=IF(OR({e},{smq}!${SM_CODE}{r}=""),"",ROUND({dsum}-({tr}),2))'   # 重复登记的行（_款式月 A 空）不算
        # W REL：停产 → 最后一个有交货双数的月份＋1（一双没交过＝1）；否则 13
        dq12 = f"{smq}!{sm_blk('双数', r)}"
        last = f'SUMPRODUCT(MAX(ISNUMBER({dq12})*({dq12}>0)*{{1,2,3,4,5,6,7,8,9,10,11,12}}))'
        ws[f'{ST_REL}{r}'] = f'=IF({e},13,IF(TRIM({ST_STAT}{r}&"")="停产",{last}+1,13))'
        for c in (ST_WM, ST_WL, ST_KEY, ST_REL):
            ws[f'{c}{r}'].font = F_HELP
    cols = [ST_SEQ, ST_CODE, ST_NAME, ST_KIND, ST_PRICE, ST_KM, ST_KL, ST_SDM, ST_SDL, ST_STAT, ST_NOTE,
            ST_OQ, ST_DQ, ST_UQ, ST_REV, ST_COST, ST_UNIT, ST_GM, ST_WIP]
    style_rows(ws, ST_R0, ST_R1, cols, auto=[ST_SEQ, ST_OQ, ST_DQ, ST_UQ, ST_REV, ST_COST, ST_UNIT, ST_GM, ST_WIP],
               fmts={ST_PRICE: MONEY, ST_KM: '0.00', ST_KL: '0.00', ST_SDM: MONEY, ST_SDL: MONEY, ST_OQ: INT, ST_DQ: INT, ST_UQ: INT,
                     ST_REV: MONEY, ST_COST: MONEY, ST_UNIT: MONEY, ST_GM: PCT, ST_CODE: '@', ST_WIP: MONEY},
               aligns={ST_NAME: AL, ST_NOTE: AL, ST_PRICE: AR, ST_REV: AR, ST_COST: AR, ST_UNIT: AR, ST_WIP: AR},
               fills={c: FILL_IN for c in cols[1:11]})
    for r in range(ST_R0, ST_R1 + 1):
        for c in (ST_WM, ST_WL, ST_KEY, ST_REL):
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, ST_WM, ST_WL, ST_KEY, ST_REL)
    dv_list(ws, f'{ST_STAT}{ST_R0}:{ST_STAT}{ST_R1}', '"在产,停产"',
            '停产＝不再做了：最后一次交货以后再记到这个款的专用成本，改按公共分摊', stop=False)
    ws.conditional_formatting.add(f'{ST_WIP}{ST_R0}:{ST_WIP}{ST_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER({ST_WIP}{ST_R0}),{ST_WIP}{ST_R0}>=0.005)'], fill=fill('FFFFEB9C'), font=F_WIP))
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True,
                        errorTitle='数字', error='填数字（空着＝1）')
    ws.add_data_validation(dv)
    dv.add(f'{ST_KM}{ST_R0}:{ST_KL}{ST_R1}')
    ws.conditional_formatting.add(f'{ST_CODE}{ST_R0}:{ST_CODE}{ST_R1}', FormulaRule(
        formula=[f'AND({ST_CODE}{ST_R0}<>"",COUNTIF(${ST_CODE}${ST_R0}:${ST_CODE}${ST_R1},{ST_CODE}{ST_R0})>1)'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(f'{ST_PRICE}{ST_R0}:{ST_PRICE}{ST_R1}', FormulaRule(
        formula=[f'AND({ST_CODE}{ST_R0}<>"",N({ST_PRICE}{ST_R0})=0)'], fill=FILL_WARN))
    ws.conditional_formatting.add(f'{ST_GM}{ST_R0}:{ST_GM}{ST_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER({ST_GM}{ST_R0}),{ST_GM}{ST_R0}<0)'], font=F_RED))
    # 右边：订单里有、档案里还没有的款式
    section(ws, ST_HDR - 1, ST_NSEQ, ST_NQTY, '订单里有、这里还没登记的款式', 'FFC00000')
    for col, t in ((ST_NSEQ, '序号'), (ST_NCODE, '款式编码'), (ST_NORD, '首次出现\n的订单'), (ST_NQTY, '订单双数')):
        put(ws, f'{col}{ST_HDR}', t, F_HDR, fill('FFC00000'), align=ACW)
    for k in range(1, ST_NEW_N + 1):
        r = ST_R0 + k - 1
        idx = f'IFERROR(SMALL({odr(OD_SNEW)},{k}),0)'
        ws[f'{ST_NSEQ}{r}'] = f'=IF({idx}=0,"",{k})'
        ws[f'{ST_NCODE}{r}'] = f'=IF({ST_NSEQ}{r}="","",INDEX({odr(OD_STY)},{idx}))'
        ws[f'{ST_NORD}{r}'] = f'=IF({ST_NSEQ}{r}="","",INDEX({odr(OD_NO)},{idx})&"")'
        ws[f'{ST_NQTY}{r}'] = f'=IF({ST_NSEQ}{r}="","",SUMIFS({odr(OD_OQ)},{odr(OD_STY)},{ST_NCODE}{r}))'
    style_rows(ws, ST_R0, ST_R0 + ST_NEW_N - 1, [ST_NSEQ, ST_NCODE, ST_NORD, ST_NQTY], auto=[ST_NSEQ, ST_NCODE, ST_NORD, ST_NQTY],
               fmts={ST_NQTY: INT})
    ws.auto_filter.ref = f'A{ST_HDR}:{ST_WIP}{ST_R1}'
    ws.freeze_panes = f'C{ST_R0}'


# ═══════════════════════════ 品名档案 ═══════════════════════════
def build_mat(wb, ctx):
    ws = wb[SH_MAT]
    widths(ws, {'A': 5, 'B': 20, 'C': 11, 'D': 6, 'E': 18, 'F': 10, 'G': 18, 'H': 10, 'I': 12, 'J': 10, 'K': 10, 'L': 2,
                'O': 5, 'P': 20, 'Q': 6, 'R': 12})
    title(ws, '品 名 档 案（材料按品名跟踪 · 可以不填）', 'K', C_ARC,
          '💡 供应商送货单上写的品名，在这里登记一下「属于哪类材料」，【材料汇总】就能按类看（面料、鞋底、包装……花了多少）。'
          '不登记也能用：没登记的品名算「未分类」，右边会自动列出来，照着抄过来选个类别就行。品名要跟送货单上写的一样（空格、括号全角半角不影响）。'
          '右边几列自动：本年数量、金额、平均单价、最后一次单价——看哪个材料涨价了。')
    ws.row_dimensions[2].height = 48
    put(ws, 'A3', '品名数', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('A3:B3')
    put(ws, 'C3', f'=SUMPRODUCT(--(TRIM({rng(SH_MAT, MT_NAME, MT_R0, MT_R1)}&"")<>""))', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    put(ws, 'D3', '没登记', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('D3:E3')
    put(ws, 'F3', f'=COUNT({dnr(DN_NNEW)})', F_KPI_V, fill('FFD9E1F2'), INT, AC)
    heads = [(MT_SEQ, '序号'), (MT_NAME, '品名及规格'), (MT_CAT, '材料类别'), (MT_UNIT, '单位'), (MT_SUP, '常用供应商'),
             (MT_REFP, '参考单价'), (MT_NOTE, '备注'), (MT_QTY, '本年数量'), (MT_AMT, '本年金额'), (MT_AVG, '平均单价'),
             (MT_LAST, '最后一次\n单价')]
    for col, t in heads:
        c = 'FF595959' if CI(col) <= CI(MT_NOTE) else 'FF548235'
        put(ws, f'{col}{MT_HDR}', t, F_HDR, fill(c), align=ACW)
    ws.row_dimensions[MT_HDR].height = 32
    ws[f'{MT_KEY}{MT_HDR}'] = '品名规范写法'
    ws[f'{MT_KEY}{MT_HDR}'].font = F_HELP
    rows = ctx.get(SH_MAT, [])
    for i, r in enumerate(range(MT_R0, MT_R1 + 1)):
        if i < len(rows):
            for c, v in rows[i].items():
                if v is not None:
                    ws[f'{c}{r}'] = v
        nm = f'{MT_NAME}{r}'
        ws[f'{MT_SEQ}{r}'] = f'=IF({nm}="","",ROW()-{MT_HDR})'
        ws[f'{MT_KEY}{r}'] = f'=IF(TRIM({nm}&"")="","",{norm(nm)})'
        ws[f'{MT_KEY}{r}'].font = F_HELP
        k = f'{MT_KEY}{r}'
        ws[f'{MT_QTY}{r}'] = f'=IF({k}="","",SUMIFS({dnr(DN_QTY)},{dnr(DN_NAMEK)},{esc(k)},{dnr(DN_OK)},1))'
        ws[f'{MT_AMT}{r}'] = f'=IF({k}="","",SUMIFS({dnr(DN_AMT)},{dnr(DN_NAMEK)},{esc(k)},{dnr(DN_OK)},1))'
        ws[f'{MT_AVG}{r}'] = f'=IF({k}="","",IF(N({MT_QTY}{r})=0,"",ROUND({MT_AMT}{r}/{MT_QTY}{r},2)))'
        last = (f'SUMPRODUCT(MAX(({dnr(DN_NAMEK)}={k})*ISNUMBER({dnr(DN_PRICE)})*({dnr(DN_PRICE)}>0)*({dnr(DN_OK)}=1)'
                f'*(ROW({dnr(DN_NAMEK)})-{DN_R0 - 1})))')
        ws[f'{MT_LAST}{r}'] = f'=IF({k}="","",IF({last}=0,"",INDEX({dnr(DN_PRICE)},{last})))'
    cols = [MT_SEQ, MT_NAME, MT_CAT, MT_UNIT, MT_SUP, MT_REFP, MT_NOTE, MT_QTY, MT_AMT, MT_AVG, MT_LAST]
    style_rows(ws, MT_R0, MT_R1, cols, auto=[MT_SEQ, MT_QTY, MT_AMT, MT_AVG, MT_LAST],
               fmts={MT_REFP: MONEY, MT_QTY: '#,##0.##;[Red]-#,##0.##;"-"', MT_AMT: MONEY, MT_AVG: MONEY, MT_LAST: MONEY},
               aligns={MT_NAME: AL, MT_SUP: AL, MT_NOTE: AL, MT_AMT: AR},
               fills={c: FILL_IN for c in cols[1:7]})
    for r in range(MT_R0, MT_R1 + 1):
        ws[f'{MT_KEY}{r}'].font = F_HELP
    hide(ws, MT_KEY)
    dv_list(ws, f'{MT_CAT}{MT_R0}:{MT_CAT}{MT_R1}', f'={MC_NAMES}', '材料类别（【基础资料】⑥）')
    dv_list(ws, f'{MT_SUP}{MT_R0}:{MT_SUP}{MT_R1}', f'={UN_NAMES}', None, stop=False)
    ws.conditional_formatting.add(f'{MT_NAME}{MT_R0}:{MT_NAME}{MT_R1}', FormulaRule(
        formula=[f'AND({MT_KEY}{MT_R0}<>"",COUNTIF(${MT_KEY}${MT_R0}:${MT_KEY}${MT_R1},{MT_KEY}{MT_R0})>1)'], fill=FILL_WARN, font=F_RED))
    section(ws, MT_HDR - 1, MT_NSEQ, MT_NAMT, '送货单里有、这里还没登记的品名', 'FFC00000')
    for col, t in ((MT_NSEQ, '序号'), (MT_NNAME, '品名'), (MT_NUNIT, '单位'), (MT_NAMT, '本年金额')):
        put(ws, f'{col}{MT_HDR}', t, F_HDR, fill('FFC00000'), align=ACW)
    for k in range(1, MT_NEW_N + 1):
        r = MT_R0 + k - 1
        idx = f'IFERROR(SMALL({dnr(DN_NNEW)},{k}),0)'
        ws[f'{MT_NSEQ}{r}'] = f'=IF({idx}=0,"",{k})'
        ws[f'{MT_NNAME}{r}'] = f'=IF({MT_NSEQ}{r}="","",INDEX({dnr(DN_NAME)},{idx})&"")'
        ws[f'{MT_NUNIT}{r}'] = f'=IF({MT_NSEQ}{r}="","",INDEX({dnr(DN_UNIT)},{idx})&"")'
        nk = esc(f'INDEX({dnr(DN_NAMEK)},{idx})')
        ws[f'{MT_NAMT}{r}'] = f'=IF({MT_NSEQ}{r}="","",SUMIFS({dnr(DN_AMT)},{dnr(DN_NAMEK)},{nk},{dnr(DN_OK)},1))'
    style_rows(ws, MT_R0, MT_R0 + MT_NEW_N - 1, [MT_NSEQ, MT_NNAME, MT_NUNIT, MT_NAMT], auto=[MT_NSEQ, MT_NNAME, MT_NUNIT, MT_NAMT],
               fmts={MT_NAMT: MONEY}, aligns={MT_NNAME: AL})
    ws.auto_filter.ref = f'A{MT_HDR}:{MT_LAST}{MT_R1}'
    ws.freeze_panes = f'C{MT_R0}'


# ═══════════════════════════ 会计科目表 ═══════════════════════════
COA_CHK = 'K'           # 检查（看得见）：自己加的科目加得对不对
COA_OWN0 = 'M'          # 隐藏：本科目自己的年初（正常方向为正；预置的「下级合计」科目 1002/1012/4001/5401 写 0，免得跟下级重复算）
COA_OWN0S = rng(SH_COA, COA_OWN0, COA_R0, COA_R1)


def _sys_codes():
    """预置科目里系统会直接记账（或下级是系统固定的）的编码：不能在它下面加下级。只有 1123、1801、2203、2221、5051、5402 不在里面。"""
    used = {c for _n, c, *_r in CATS if not c.startswith('@')}
    used |= {c for c, _cf in FE_CLASSES.values()} | {v[1] for v in DP_TYPES.values()} | set(FA_USES.values())
    used |= set(COMP_CODE.values()) | set(COGS_CODE.values()) | {MOH, '1122', '5001', '2202', '2211', '1602', '1001'}
    return [c for c, _n, _cl, _d, up, _l, src in COA if c in used or src in ('sum', 'acct') or up in ('1002', '1012', '4001', '5401')]


AX_SYS_COL, AX_SYS0 = 'F', 2      # 【_辅助】F 列：上面这个清单（F2 起）
AX_SYS = rng(SH_AUX, AX_SYS_COL, AX_SYS0, AX_SYS0 + len(COA) - 1)


def build_coa(wb, ctx):
    ws = wb[SH_COA]
    widths(ws, {'A': 10, 'B': 26, 'C': 7, 'D': 7, 'E': 9, 'F': 16, 'G': 13, 'H': 13, 'I': 6, 'J': 40, COA_CHK: 40})
    sys_txt = '1001/1002xx/1012xx 资金、1122/2202/1221/2241 往来、2211、1601/1602、400101～400104、4101、5001、5401/540101～540104、5601/5602/5603 费用等'
    title(ws, '会 计 科 目 表（小企业会计准则 · 期初余额）', COA_CHK, C_ARC,
          '💡 平时不用动。「年初余额（手填）」只填淡黄格（应付职工薪酬、实收资本、年初未分配利润、借款……建账日那天的数）；'
          '资金账户、往来单位、固定资产的期初在各自的表里填，这里自动取（灰格）。正数＝科目正常方向的余额（资产在借方、负债权益在贷方）。'
          '第 3 行「期初平不平」要是 √ 才对（每个科目只算它自己的年初，上级不重复算）。'
          '要加科目：在下面空行加；下级科目编码用上级开头（比如 2221 应交税费下面加 222101），「上级编码」要填，年初余额填在下级上，报表里上级自动包含下级。'
          f'系统自动记账的科目（{sys_txt}）不要加下级：系统分录照样记在原科目上，下级对不上；预置科目里只有 1123、1801、2203、2221、5051、5402 能加下级。'
          '费用要分细用【基础资料】④ 费用项目。K 列提示加得对不对。')
    ws.row_dimensions[2].height = max(ws.row_dimensions[2].height or 0, 80)
    put(ws, 'A3', '期初平不平', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('A3:B3')
    dr, own = rng(SH_COA, COA_DIR, COA_R0, COA_R1), COA_OWN0S
    s_d, s_c = f'SUMIFS({own},{dr},"借")', f'SUMIFS({own},{dr},"贷")'
    diff = f'ROUND({s_d}-{s_c},2)'
    put(ws, 'C3', f'=IF({diff}=0,"√ 平",IF({diff}>0,"✗ 借方比贷方多 ","✗ 贷方比借方多 ")&TEXT(ABS({diff}),"#,##0.00"))',
        F_KPI_V, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('C3:F3')
    put(ws, 'G3', '借方合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'H3', f'=ROUND({s_d},2)', F_KPI_V, fill('FFD9E1F2'), MONEY, AC)
    put(ws, 'I3', f'="贷方合计　"&TEXT({s_c},"#,##0.00")', F_KPI_V, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('I3:J3')
    ck = f'${COA_CHK}${COA_R0}:${COA_CHK}${COA_R1}'
    put(ws, f'{COA_CHK}3', f'=IF(COUNTIF({ck},"✗*")>0,"✗ 有 "&COUNTIF({ck},"✗*")&" 个科目加得不对：看下面 K 列",'
                           f'IF(COUNTIF({ck},"⚠*")>0,"⚠ 有 "&COUNTIF({ck},"⚠*")&" 个科目要看看：K 列",""))', F_TXTB, fill('FFD9E1F2'), align=AL)
    heads = [(COA_CODE, '科目编码'), (COA_NAME, '科目名称'), (COA_CLS, '类别'), (COA_DIR, '方向'), (COA_UP, '上级编码'),
             (COA_LINE, '报表项目'), (COA_OPEN_IN, '年初余额\n（手填）'), (COA_OPEN, '年初余额\n（用的）'), (COA_LEAF, '末级'), (COA_NOTE, '说明'),
             (COA_CHK, '检查（自己加的科目）')]
    header(ws, COA_HDR, heads, 'FF595959')
    ws[f'{COA_OWN0}{COA_HDR}'] = '本科目自己的年初\n（下级合计行＝0）'
    ws[f'{COA_OWN0}{COA_HDR}'].font = F_HELP
    opens = ctx.get('coa_open', {})
    rows = {}
    for i, (code, name, cls, d, up, line, src) in enumerate(COA):
        rows[code] = COA_R0 + i
    notes = {'1601': '建账日以前买的固定资产原值（【基础资料】⑦）', '1602': '建账前已提折旧合计（【基础资料】⑦）',
             '1122': '【往来单位】期初应收合计', '2202': '【往来单位】期初应付合计', '1221': '【往来单位】期初其他应收合计',
             '2241': '【往来单位】期初其他应付合计', '1001': '【基础资料】② 现金账户期初', '4001': '建账日车间里没做完的（在制品），一般填 0',
             '3104': '年初未分配利润（期初资产−负债−实收资本，倒挤也行：C3 显示「借方比贷方多 x」就把这里加上 x，「贷方比借方多 x」就减掉 x）',
             '2211': '建账日还没发的工资',
             '4101': '月末自动转到生产成本-制造费用，余额应为 0', '400104': '月末从制造费用自动转进来（不要手工记）',
             '5401': '月末按【成本分摊表】自动结转（不要手工记 540101～540104）', '1123': '一般不用：预付供应商的钱直接记应付账款（资产负债表自动按余额方向拆成预付）'}
    for i, (code, name, cls, d, up, line, src) in enumerate(COA):
        r = COA_R0 + i
        ws[f'{COA_CODE}{r}'] = code
        if name.startswith('@bank') or name.startswith('@pay'):
            pre = '银行存款-' if name.startswith('@bank') else '其他货币资金-'
            ws[f'{COA_NAME}{r}'] = f'=IFERROR("{pre}"&INDEX({AC_NAMES},MATCH({COA_CODE}{r},{AC_CODES},0)),"{pre}（备用）")'
        else:
            ws[f'{COA_NAME}{r}'] = name
        ws[f'{COA_CLS}{r}'], ws[f'{COA_DIR}{r}'], ws[f'{COA_UP}{r}'], ws[f'{COA_LINE}{r}'] = cls, d, up or None, line
        h = f'{COA_OPEN}{r}'
        if src == 'acct':
            ws[h] = f'=SUMIF({AC_CODES},{COA_CODE}{r},{AC_OPENS})'
        elif src == 'sum':
            kids = [rows[c] for c, *_x in COA if _x[3] == code]
            ws[h] = f'=SUM({COA_OPEN}{min(kids)}:{COA_OPEN}{max(kids)})'
        elif src.startswith('unit:'):
            col = {'AR': UN_AR0, 'AP': UN_AP0, 'OR': UN_OR0, 'OP': UN_OP0}[src[5:]]
            ws[h] = f'=SUM({unr(col)})'
        elif src == 'fa':
            fd, fc = rng(SH_BASE, FA_DATE, FA_R0, FA_R1), rng(SH_BASE, FA_COST, FA_R0, FA_R1)
            ws[h] = f'=SUMPRODUCT(ISNUMBER({fd})*({fd}<{P["OPEN"]}),{fc})'
        elif src == 'dep':
            fd, f0 = rng(SH_BASE, FA_DATE, FA_R0, FA_R1), rng(SH_BASE, FA_DEP0, FA_R0, FA_R1)
            ws[h] = f'=SUMPRODUCT(ISNUMBER({fd})*({fd}<{P["OPEN"]}),{f0})'
        else:
            ws[h] = f'=N({COA_OPEN_IN}{r})'
            if code in opens:
                ws[f'{COA_OPEN_IN}{r}'] = opens[code]
        if code in notes:
            ws[f'{COA_NOTE}{r}'] = notes[code]
            if len(notes[code]) > 20:            # 说明列 40 宽，一行 20 个字左右
                _fit_h(ws, r, notes[code], 20, 14)
    for r in range(COA_R0 + len(COA), COA_R1 + 1):
        ws[f'{COA_OPEN}{r}'] = f'=IF({COA_CODE}{r}="","",N({COA_OPEN_IN}{r}))'
    for r in range(COA_R0, COA_R1 + 1):
        ws[f'{COA_LEAF}{r}'] = f'=IF({COA_CODE}{r}="","",IF(COUNTIF(${COA_UP}${COA_R0}:${COA_UP}${COA_R1},{COA_CODE}{r}&"")=0,1,0))'
    # 本科目自己的年初（期初平不平、科目余额表核对都按这一列；预置的「下级合计」行写 0）
    codes = f'${COA_CODE}${COA_R0}:${COA_CODE}${COA_R1}'
    for i in range(COA_R1 - COA_R0 + 1):
        r = COA_R0 + i
        if i < len(COA) and COA[i][6] == 'sum':
            ws[f'{COA_OWN0}{r}'] = '=0'
        else:
            ws[f'{COA_OWN0}{r}'] = f'=IF(TRIM({COA_CODE}{r}&"")="",0,N({COA_OPEN}{r}))'
        ws[f'{COA_OWN0}{r}'].font = F_HELP
        dup = f'COUNTIF({codes},TRIM({COA_CODE}{r}&""))>1'
        if i < len(COA):
            ws[f'{COA_CHK}{r}'] = f'=IF({dup},"✗ 编码跟下面加的科目重复：把下面那个换一个编码","")'
            continue
        c, u = f'TRIM({COA_CODE}{r}&"")', f'TRIM({COA_UP}{r}&"")'
        tc = f'TRIM({codes}&"")'
        sys_pre = f'SUMPRODUCT(({AX_SYS}<>"")*(LEN({AX_SYS})<LEN({c}))*(LEFT({c},LEN({AX_SYS}))={AX_SYS}))'
        sys_up = f'IF({u}="",0,COUNTIF({AX_SYS},{u})+COUNTIF({CT_OPPS},{u})+COUNTIF({FE_CODES},{u}))'
        pre = f'SUMPRODUCT(({tc}<>"")*(LEN({tc})<LEN({c}))*(LEFT({c},LEN({tc}))={tc}))'
        d = f'TRIM({COA_DIR}{r}&"")'
        ws[f'{COA_CHK}{r}'] = (
            f'=IF({c}="","",IF({dup},"✗ 编码重复：换一个",'
            f'IF({sys_pre}+{sys_up}>0,"✗ 不能加在系统记账的科目下面：上级编码空着，换个新的一级编码",'
            f'IF(AND({u}<>"",COUNTIF({codes},{u})=0),"✗ 上级编码 "&{u}&" 本表里没有",'
            f'IF(AND({u}<>"",LEFT({c},LEN({u}))<>{u}),"⚠ 下级编码要用上级编码开头（如 "&{u}&"01）",'
            f'IF(AND({u}="",{pre}>0),"⚠ 编码以已有科目开头：「上级编码」要填那个科目",'
            f'IF(AND(N({COA_OPEN}{r})<>0,{d}<>"借",{d}<>"贷"),"✗ 有年初余额，要选方向（借 / 贷）",'
            f'IF(TRIM({COA_CLS}{r}&"")="","⚠ 选类别（资产 / 负债 / 权益 / 成本 / 损益）","√"))))))))')
    cols = [COA_CODE, COA_NAME, COA_CLS, COA_DIR, COA_UP, COA_LINE, COA_OPEN_IN, COA_OPEN, COA_LEAF, COA_NOTE]
    style_rows(ws, COA_R0, COA_R1, cols, auto=[COA_OPEN, COA_LEAF],
               fmts={COA_CODE: '@', COA_UP: '@', COA_OPEN_IN: MONEY, COA_OPEN: MONEY},
               aligns={COA_CODE: AL, COA_NAME: AL, COA_NOTE: ALW, COA_OPEN_IN: AR, COA_OPEN: AR, COA_LINE: AL})
    for i, (code, name, cls, d, up, line, src) in enumerate(COA):
        r = COA_R0 + i
        if src == 'in':
            ws[f'{COA_OPEN_IN}{r}'].fill = FILL_IN
        else:
            ws[f'{COA_OPEN_IN}{r}'].fill = FILL_AUTO
            ws[f'{COA_OPEN_IN}{r}'].value = None
        if not up:
            for c in (COA_CODE, COA_NAME):
                ws[f'{c}{r}'].font = F_TXTB
        for c in (COA_CODE, COA_NAME, COA_CLS, COA_DIR, COA_UP, COA_LINE):
            ws[f'{c}{r}'].font = F_TXTB if not up and c in (COA_CODE, COA_NAME) else F_TXT
    for r in range(COA_R0 + len(COA), COA_R1 + 1):
        for c in (COA_CODE, COA_NAME, COA_CLS, COA_DIR, COA_UP, COA_LINE, COA_OPEN_IN, COA_NOTE):
            ws[f'{c}{r}'].fill = FILL_IN
    dv_list(ws, f'{COA_CLS}{COA_R0}:{COA_CLS}{COA_R1}', '"资产,负债,权益,成本,损益"', None)
    dv_list(ws, f'{COA_DIR}{COA_R0}:{COA_DIR}{COA_R1}', '"借,贷"', None)
    style_rows(ws, COA_R0, COA_R1, [COA_CHK], auto=[COA_CHK], aligns={COA_CHK: ALW})
    for r in range(COA_R0, COA_R1 + 1):
        ws[f'{COA_CHK}{r}'].font = Font(name=YH, sz=9, color='FF404040')
    hide(ws, COA_OWN0)
    ws.conditional_formatting.add('C3', FormulaRule(formula=['LEFT($C$3,1)="✗"'], fill=FILL_WARN, font=F_RED))
    k0 = f'{COA_CHK}{COA_R0}'
    for rg in (f'{COA_CHK}{COA_R0}:{COA_CHK}{COA_R1}', f'{COA_CHK}3'):
        a = rg.split(':')[0]
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({a},1)="✗"'], fill=FILL_WARN, font=F_RED))
        ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({a},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(f'{COA_CHK}{COA_R0}:{COA_CHK}{COA_R1}', FormulaRule(
        formula=[f'LEFT({k0},1)="√"'], font=Font(name=YH, sz=9, color='FF548235')))
    ws.freeze_panes = f'C{COA_R0}'


# ═══════════════════════════ _辅助 ═══════════════════════════
def build_aux(wb, ctx):
    ws = wb[SH_AUX]
    ws['A1'] = f'=COUNT({unr(UN_ARK)})'
    for k in range(1, JE_NCUST + 1):
        ws[f'A{AX_CUST0 + k - 1}'] = f'=IFERROR(INDEX({UN_NAMES},SMALL({unr(UN_ARK)},{k}))&"","")'
    ws['B1'] = '客户清单（【往来单位】类型＝客户，按顺序，最多 %d 个；记账分录「收入」块按它排）' % JE_NCUST
    ws['C1'], ws['D1'] = '月', '月份名'
    for m in range(1, 13):
        ws[f'C{m + 1}'] = m
        ws[f'D{m + 1}'] = f'{m}月'
    ws[f'{AX_SYS_COL}1'] = '系统直接记账的科目（【会计科目表】不能在它们下面加下级）'
    for i, c in enumerate(_sys_codes()):
        ws[f'{AX_SYS_COL}{AX_SYS0 + i}'] = c
        ws[f'{AX_SYS_COL}{AX_SYS0 + i}'].number_format = '@'


def build(wb, ctx):
    build_base(wb, ctx)
    build_unit(wb, ctx)
    build_style(wb, ctx)
    build_mat(wb, ctx)
    build_coa(wb, ctx)
    build_aux(wb, ctx)
