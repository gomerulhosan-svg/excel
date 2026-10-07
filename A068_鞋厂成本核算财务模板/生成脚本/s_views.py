# -*- coding: utf-8 -*-
"""查看表 8 张（第二阶段）：【材料汇总】【晚到单据】【应付账款汇总】【供应商对账单】【客户对账单】【费用汇总】【工资汇总】【账户余额表】。

   · 全是公式、只读（make 会整张保护），顶上的选择格亮黄（不锁）。
   · 被别的表引用的格子只有两张对账单的选择格（SS_UNIT/SS_D1/SS_D2、CS_UNIT/CS_D1/CS_D2，位置按 layout.py）。
   · 各表右边隐藏的列是本表自己用的帮手列（第 k 个是源表第几行、排序键、行的种类……）。
   · 清单一律「★序号键 / 排序键 + SMALL(k) / LARGE(k)」取第 k 个，k 超过个数就不算（IF 短路），取不到显示空。
     INDEX(区域, k) 的 k 可能是 0 的，一律放在 IF(k=0, …) 的另一支里（普通格子里 INDEX(区域,0) 会按行号取交叉，超出区域就 #VALUE!）。
   · 两张对账单的「合计、签字栏」是浮动的：紧跟在最后一笔下面（不会在 300 行以后才出现），边框用条件格式只画有内容的行。
   · 月份选择格出厂空着＝最新月份（【基础资料】LASTM；还没业务＝建账月），公式里用隐藏的「实际用的月份」帮手格。
   · 不按往来单位、不按费用项目的科目发生额用 mat()（科目余额表隐藏的科目×月矩阵），不再扫记账分录。
"""
import datetime as _dt
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Border
from common import *
from layout import *          # 注意：layout 的 C_IN / C_RPT 等颜色覆盖 common 的同名常量

H_VIEW = 'FF548235'                     # 查看表表头：绿（跟 s_cost 一样）
H_DARK = 'FF375623'                     # 区块标题
KPI_FILL = fill('FFD9E1F2')
F_VAL = Font(name=YH, sz=11, bold=True, color='FF1F3864')
F_GOOD = Font(name=YH, sz=10, bold=True, color='FF00B050')
F_REDN = Font(name=YH, sz=10, color='FFC00000')
F_BOLD = Font(bold=True)
FILL_YEL = fill('FFFFEB9C')
CHKF = '"✗ 差 "#,##0.00;"✗ 差 "-#,##0.00;"√"'      # 核对格：0 显示 √
MHDR = '0"月"'
DAYF = '0.0;-0.0;"-"'
QTYF = 'General'
CNTF = '0;-0;"-"'
BD_CF = Border(left=thin, right=thin, top=thin, bottom=thin)

YR, OPEN, LOCK, LAGD, CO = (P[k] for k in ('YEAR', 'OPEN', 'LOCK', 'LAGD', 'CO'))   # LOCK 只是显示（已结账到几月）
XFER = '内部转账'
YEAR0 = next(v for k, _l, v, *_r in PARAMS if k == 'YEAR')          # 选择格的常量默认日期用（2026）
COA_OPEN_OF = lambda code: f'SUMIF({COA_CODES},"{code}",{COA_OPENS})'
# 建账月（建账日期在本年度＝它的月份，否则 1）；「最新月份」＝【基础资料】LASTM，还没业务（0）时＝建账月——保证落在 1～12
OPEN_M = f'IF(ISNUMBER({OPEN}),IF(YEAR({OPEN})={YR},MONTH({OPEN}),1),1)'
LATEST = f'MAX({OPEN_M},MIN(12,INT(N({P["LASTM"]}))))'


def _mx(sel, empty='latest'):
    """选择格 sel 填的是 1～12 就用它（取整），否则（空着 / 填错）用最新月份（empty='latest'）或 1（empty=1）"""
    dft = LATEST if empty == 'latest' else str(empty)
    return f'IF(AND(ISNUMBER({sel}),{sel}>=1,{sel}<13),INT({sel}),{dft})'


def _mx_ok(sel):
    return f'AND(ISNUMBER({sel}),{sel}>=1,{sel}<13)'


# ─────────────────────────── 小工具 ───────────────────────────
def _lbl(ws, coord, text, merge=None, fill_=KPI_FILL, font=F_KPI_L):
    if merge:
        ws.merge_cells(merge)
    put(ws, coord, text, font, fill_, align=ACW)


def _val(ws, coord, f, fmt=MONEY, merge=None, font=F_VAL, fill_=KPI_FILL, align=AC):
    if merge:
        ws.merge_cells(merge)
    put(ws, coord, f, font, fill_, fmt, align)


def _mhdr(ws, row, c0, color=H_VIEW, n=12):
    """表头：1～12 月写成数字 1～12（格式显示「1月」），公式里就能直接拿表头当月份用"""
    for m in range(1, n + 1):
        put(ws, f'{CL(CI(c0) + m - 1)}{row}', m, F_HDR, fill(color), MHDR, ACW)


def _help(ws, hdr_row, labels, r0, r1):
    """隐藏帮手列：表头小灰字、整列小灰字、隐藏"""
    for c, t in labels:
        if t is not None:
            ws[f'{c}{hdr_row}'] = t
            ws[f'{c}{hdr_row}'].font = F_HELP
        for r in range(r0, r1 + 1):
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in labels])


def _neg_red(ws, rg, first):
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}<0)'], font=F_REDN))


def _chk_cf(ws, rg, first):
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="✗"'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], fill=FILL_YEL))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="√"'], font=F_GOOD))


def _chk_num_cf(ws, rg, first):
    """核对格（数字，0 显示 √）：不是 0 红底"""
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({first}),ROUND({first},2)<>0)'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({first}),ROUND({first},2)=0)'], font=F_GOOD))


def _plain_rows(ws, r0, r1, cols, fmts=None, aligns=None, font=F_AUTO):
    """清单区：不画框不填色（有内容的行用条件格式画框）"""
    fmts, aligns = fmts or {}, aligns or {}
    for r in range(r0, r1 + 1):
        for c in cols:
            x = ws[f'{c}{r}']
            x.font = font
            x.alignment = aligns.get(c, AC)
            if c in fmts:
                x.number_format = fmts[c]


def _dtxt(x):
    """日期 → 「2026/9/5」文字（不用 TEXT 的日期格式，免得各地区设置不一样）"""
    return f'YEAR({x})&"/"&MONTH({x})&"/"&DAY({x})'


def _dtxt_cn(x):
    return f'YEAR({x})&"年"&MONTH({x})&"月"&DAY({x})&"日"'


def _home(ws, last):
    col = CL(CI(last) + 1)
    home_link(ws, f'{col}1')
    ws.column_dimensions[col].width = max(ws.column_dimensions[col].width or 0, 10)


def _row_total(ws, r, cols, fmt=MONEY, fill_=FILL_TOT, font=F_AUTOB):
    for c in cols:
        x = ws[f'{c}{r}']
        x.font, x.fill, x.border, x.number_format = font, fill_, BD, fmt
        x.alignment = AR


# ═══════════════════════════ 材料汇总 ═══════════════════════════
MS_HDR = 5                                    # ① 表头（1～12 月是数字）
MS_R0 = 6                                     # 第一个材料类别
NMC = MC_R1 - MC_R0 + 1                       # 20 个类别
MS_RN = MS_R0 + NMC                           # 未分类
MS_RDN = MS_RN + 1                            # 送货单小计
MS_RCASH = MS_RDN + 1                         # 现付材料（日记账）
MS_RMJ = MS_RCASH + 1                         # 手工分录
MS_RTOT = MS_RMJ + 1                          # 材料合计
MS_RAL = MS_RTOT + 1                          # 成本分摊表「材料 本月发生」
MS_RCHK = MS_RAL + 1                          # 核对
MN_SEC = MS_RCHK + 2                          # ② 品名清单
MN_HDR = MN_SEC + 1
MN_R0 = MN_HDR + 1
MN_N = 400
MN_R1 = MN_R0 + MN_N - 1
MN_IDX, MN_KEY, MN_LAST, MN_MAX = 'R', 'S', 'T', 'U'   # 隐藏：第 k 个品名在送货单第几行、品名规范写法、最后一次（日期×100000＋行）、最高单价
MS_HP, MS_HK = 'V', 'W'                       # 隐藏（跟【送货单登记】一行对一行）：能记账且有单价的单价；它的「日期×100000＋行」


def build_msum(wb, ctx):
    ws = wb[SH_MSUM]
    widths(ws, {'A': 6, 'B': 22, **{CL(3 + i): 10.5 for i in range(12)}, 'O': 13, 'P': 13, 'Q': 2})
    title(ws, '材 料 汇 总（按材料类别 × 月 · 按品名看数量、金额、单价）', 'P', C_VIEW,
          '💡 全自动。材料按「成本月份」算（就是送货单上的送货日期那个月；【基础资料】⑩「每月结账日期」填了的月份，'
          f'结账以后才收到的那个月的送货单，算到收单那个月，已结账月份的数不再变——在【{SH_LATE}】能看到挪了哪些）。'
          '① 每个材料类别每个月花了多少钱：送货单按【品名档案】登记的类别分，没登记类别的品名算「未分类」（去【品名档案】登记一下就分进去了）；'
          '另外加上资金日记账里「现付材料」（没有送货单、当场付钱的）和手工分录记到材料成本的。最下面核对：材料合计应该等于【成本分摊表】的「材料 本月发生」。'
          '② 每个品名（同一个品名不同颜色合在一起）：所选月份和全年的数量、金额、平均单价，最高 / 最低 / 最后一次的单价——看哪个材料涨价了。'
          'C3、E3 选起止月份（黄格）：起始月空着＝1 月，截止月空着＝最新月份（有业务的最后一个月）。', h2=64)
    _home(ws, 'P')
    selector(ws, 'B3', '起始月', 'C3', None, f'={AX_M}', MHDR, '从几月（1～12）；空着＝1 月')
    selector(ws, 'D3', '截止月', 'E3', None, f'={AX_M}', MHDR, '到几月（1～12）；空着＝最新月份')
    # 实际用的起止月（隐藏格）：C3 空＝1，E3 空＝最新月份
    m1, m2 = f'${MN_IDX}$3', f'${MN_KEY}$3'
    ws[f'{MN_IDX}3'] = f'={_mx("$C$3", 1)}'
    ws[f'{MN_KEY}3'] = f'={_mx("$E$3")}'
    ws[f'{MN_IDX}3'].font = ws[f'{MN_KEY}3'].font = F_HELP
    _lbl(ws, 'F3', (f'={m1}&"～"&{m2}&" 月材料合计"&IF({_mx_ok("$E$3")},"",CHAR(10)&'
                    f'IF(TRIM($E$3&"")="","（空着＝最新月份）","⚠ 截止月要选 1～12"))'), 'F3:G3')
    _val(ws, 'H3', f'=O{MS_RTOT}', MONEY, 'H3:I3')
    _lbl(ws, 'J3', '全年\n材料合计', 'J3:K3')
    _val(ws, 'L3', f'=P{MS_RTOT}', MONEY, 'L3:M3')
    put(ws, 'N3', f'=IF({m1}>{m2},"⚠ 起始月比截止月大",IF(ROUND(P{MS_RCHK},2)<>0,"✗ 跟成本分摊表对不上（看第 {MS_RCHK} 行）","√ 跟成本分摊表对上了"))',
        F_KPI_L, KPI_FILL, align=ACW)
    ws.merge_cells('N3:P3')
    _chk_cf(ws, 'N3', '$N$3')
    ws.row_dimensions[3].height = 32

    # ── ① 类别 × 月
    section(ws, MS_HDR - 1, 'A', 'P', '① 按材料类别 × 月（元；送货单按成本月份，能记账的行）', H_DARK)
    put(ws, f'A{MS_HDR}', '序号', F_HDR, fill(H_VIEW), align=ACW)
    put(ws, f'B{MS_HDR}', '材料类别', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, MS_HDR, 'C')
    put(ws, f'O{MS_HDR}', '="所选月份"&CHAR(10)&"（"&' + m1 + '&"～"&' + m2 + '&"月）"', F_HDR, fill(H_VIEW), align=ACW)
    put(ws, f'P{MS_HDR}', '全年', F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[MS_HDR].height = 34
    MC = [CL(3 + i) for i in range(12)]
    hdr_m = f'$C${MS_HDR}:$N${MS_HDR}'

    def period(r):
        return f'SUMIFS(C{r}:N{r},{hdr_m},">="&{m1},{hdr_m},"<="&{m2})'

    for k in range(NMC):
        r = MS_R0 + k
        b = f'$B{r}'
        ws[f'A{r}'] = f'=IF({b}="","",{k + 1})'
        ws[f'B{r}'] = f"={q(SH_BASE)}!${MC_NAME}${MC_R0 + k}&\"\""
        for c in MC:
            ws[f'{c}{r}'] = (f'=IF({b}="","",SUMIFS({dnr(DN_AMT)},{dnr(DN_CAT)},{esc(b)},{dnr(DN_CM)},{c}${MS_HDR},'
                             f'{dnr(DN_OK)},1))')
        ws[f'O{r}'] = f'=IF({b}="","",{period(r)})'
        ws[f'P{r}'] = f'=IF({b}="","",SUM(C{r}:N{r}))'
    style_rows(ws, MS_R0, MS_R0 + NMC - 1, ['A', 'B'] + MC + ['O', 'P'], auto=['A', 'B'] + MC + ['O', 'P'],
               fmts={c: MONEY for c in MC + ['O', 'P']}, aligns={'B': AL, **{c: AR for c in MC + ['O', 'P']}})
    for r in range(MS_R0, MS_R0 + NMC):
        for c in ['A', 'B'] + MC + ['O', 'P']:
            ws[f'{c}{r}'].fill = FILL_NONE
    # 未分类、送货单小计、现付、手工、合计、成本分摊表、核对
    special = [
        (MS_RN, '未分类（品名没在【品名档案】登记类别的）', lambda c: f'=ROUND({c}{MS_RDN}-SUM({c}{MS_R0}:{c}{MS_RN - 1}),2)', None),
        (MS_RDN, '送货单小计', lambda c: f'=SUMIFS({dnr(DN_AMT)},{dnr(DN_CM)},{c}${MS_HDR},{dnr(DN_OK)},1)', FILL_SUB),
        (MS_RCASH, '现付材料（资金日记账）', lambda c: f'=SUMIFS({jr(J_CAMT)},{jr(J_COMP)},"材料",{jr(J_CM)},{c}${MS_HDR})', None),
        (MS_RMJ, '手工分录（记到材料成本的）', lambda c: f'=SUMIFS({mjr(MJ_CAMT)},{mjr(MJ_COMP)},"材料",{mjr(MJ_CM)},{c}${MS_HDR})', None),
        (MS_RTOT, '材料合计', lambda c: f'=ROUND({c}{MS_RDN}+{c}{MS_RCASH}+{c}{MS_RMJ},2)', FILL_TOT),
    ]
    for r, label, fm, fl in special:
        put(ws, f'A{r}', None, F_TXTB, fl)
        put(ws, f'B{r}', label, F_TXTB if fl else F_TXT, fl, align=AL)
        for c in MC:
            put(ws, f'{c}{r}', fm(c), F_AUTOB if fl else F_AUTO, fl, MONEY, AR)
        put(ws, f'O{r}', f'={period(r)}', F_AUTOB if fl else F_AUTO, fl, MONEY, AR)
        put(ws, f'P{r}', f'=SUM(C{r}:N{r})', F_AUTOB if fl else F_AUTO, fl, MONEY, AR)
    r = MS_RAL
    put(ws, f'A{r}', None, F_TXT)
    put(ws, f'B{r}', '【成本分摊表】材料 本月发生', F_TXT, align=AL)
    for i, c in enumerate(MC):
        put(ws, f'{c}{r}', f"=N({al('材料', '本月发生', i + 1)})", F_AUTO, None, MONEY, AR)
    put(ws, f'O{r}', f'={period(r)}', F_AUTO, None, MONEY, AR)
    put(ws, f'P{r}', f'=SUM(C{r}:N{r})', F_AUTO, None, MONEY, AR)
    r = MS_RCHK
    put(ws, f'A{r}', None, F_TXTB, FILL_SUB)
    put(ws, f'B{r}', '核对（合计 − 成本分摊表，应为 0）', F_TXTB, FILL_SUB, align=AL)
    for c in MC + ['O', 'P']:
        put(ws, f'{c}{r}', f'=ROUND({c}{MS_RTOT}-{c}{MS_RAL},2)', F_AUTOB, FILL_SUB, CHKF, AC)
    _chk_num_cf(ws, f'C{r}:P{r}', f'C{r}')
    _neg_red(ws, f'C{MS_R0}:P{MS_RAL}', f'C{MS_R0}')

    # ── ② 品名清单
    section(ws, MN_SEC, 'A', 'L', f'② 按品名（送货单 · 同一个品名不同颜色合在一起 · 按第一次出现的顺序 · 最多 {MN_N} 个）', H_DARK)
    cnt = f'${MN_IDX}${MN_SEC}'
    ws[cnt.replace('$', '')] = f'=COUNT({dnr(DN_NFIRST)})'
    put(ws, f'M{MN_SEC}', f'=IF({cnt}>{MN_N},"⚠ 品名有 "&{cnt}&" 个，只列前 {MN_N} 个","共 "&{cnt}&" 个品名")', F_KPI_L, KPI_FILL, align=AC)
    ws.merge_cells(f'M{MN_SEC}:P{MN_SEC}')
    heads = [('A', '序号'), ('B', '品名及规格'), ('C', '材料类别'), ('D', '单位'),
             ('E', '="所选月份"&CHAR(10)&"数量"'), ('F', '="所选月份"&CHAR(10)&"金额"'), ('G', '="所选月份"&CHAR(10)&"平均单价"'),
             ('H', '全年\n数量'), ('I', '全年\n金额'), ('J', '全年\n平均单价'), ('K', '最高\n单价'), ('L', '最低\n单价'),
             ('M', '最后一次\n单价'), ('N', '最后一次\n送货日期'), ('O', '最后一次\n供应商'), ('P', '送货\n行数')]
    header(ws, MN_HDR, heads, H_VIEW)
    nk, ok, cm = dnr(DN_NAMEK), dnr(DN_OK), dnr(DN_CM)
    # 跟【送货单登记】一行对一行的数字帮手列（SUMPRODUCT(MAX()) 里不能碰到文字，所以先把单价、日期变成干净的数字）
    for r in range(DN_R0, DN_R1 + 1):
        x = lambda col: f'{q(SH_DN)}!${col}{r}'
        ws[f'{MS_HP}{r}'] = f'=IF({x(DN_OK)}=1,IF(ISNUMBER({x(DN_PRICE)}),IF({x(DN_PRICE)}>0,{x(DN_PRICE)},0),0),0)'
        ws[f'{MS_HK}{r}'] = f'=IF({MS_HP}{r}>0,INT({x(DN_DATE)})*100000+ROW()-{DN_R0 - 1},0)'
    hp, hk = f'${MS_HP}${DN_R0}:${MS_HP}${DN_R1}', f'${MS_HK}${DN_R0}:${MS_HK}${DN_R1}'
    for r in range(MN_R0, MN_R1 + 1):
        R, S, T = f'${MN_IDX}{r}', f'${MN_KEY}{r}', f'${MN_LAST}{r}'
        e = f'{R}=0'
        k = f'ROW()-{MN_R0 - 1}'
        ws[f'{MN_IDX}{r}'] = f'=IF({k}>{cnt},0,SMALL({dnr(DN_NFIRST)},{k}))'
        ws[f'{MN_KEY}{r}'] = f'=IF({e},"",INDEX({nk},{R}))'
        crit = f'{nk},{esc(S)},{ok},1'
        per = f',{cm},">="&{m1},{cm},"<="&{m2}'
        mx = f'${MN_MAX}{r}'
        ws[f'{MN_LAST}{r}'] = f'=IF({e},0,SUMPRODUCT(MAX(({nk}={S})*{hk})))'
        ws[f'{MN_MAX}{r}'] = f'=IF({e},0,SUMPRODUCT(MAX(({nk}={S})*{hp})))'
        ws[f'A{r}'] = f'=IF({e},"",{k})'
        ws[f'B{r}'] = f'=IF({e},"",INDEX({dnr(DN_NAME)},{R})&"")'
        ws[f'C{r}'] = f'=IF({e},"",INDEX({dnr(DN_CAT)},{R})&"")'
        ws[f'D{r}'] = f'=IF({e},"",INDEX({dnr(DN_UNIT)},{R})&"")'
        ws[f'E{r}'] = f'=IF({e},"",SUMIFS({dnr(DN_QTY)},{crit}{per}))'
        ws[f'F{r}'] = f'=IF({e},"",SUMIFS({dnr(DN_AMT)},{crit}{per}))'
        ws[f'G{r}'] = f'=IF({e},"",IF(N(E{r})=0,"",ROUND(F{r}/E{r},2)))'
        ws[f'H{r}'] = f'=IF({e},"",SUMIFS({dnr(DN_QTY)},{crit}))'
        ws[f'I{r}'] = f'=IF({e},"",SUMIFS({dnr(DN_AMT)},{crit}))'
        ws[f'J{r}'] = f'=IF({e},"",IF(N(H{r})=0,"",ROUND(I{r}/H{r},2)))'
        ws[f'K{r}'] = f'=IF(N({mx})=0,"",{mx})'
        ws[f'L{r}'] = f'=IF(N({mx})=0,"",ROUND(1000000-SUMPRODUCT(MAX(({nk}={S})*({hp}>0)*(1000000-{hp}))),6))'
        ws[f'M{r}'] = f'=IF({T}=0,"",INDEX({dnr(DN_PRICE)},MOD({T},100000)))'
        ws[f'N{r}'] = f'=IF({T}=0,"",INT({T}/100000))'
        ws[f'O{r}'] = f'=IF({T}=0,"",INDEX({dnr(DN_UNITK)},MOD({T},100000))&"")'
        ws[f'P{r}'] = f'=IF({e},"",COUNTIFS({crit}))'
    cols = [CL(i) for i in range(1, 17)]
    style_rows(ws, MN_R0, MN_R1, cols, auto=cols,
               fmts={'A': CNTF, 'E': QTYF, 'F': MONEY, 'G': MONEY, 'H': QTYF, 'I': MONEY, 'J': MONEY, 'K': MONEY, 'L': MONEY,
                     'M': MONEY, 'N': DATE, 'P': CNTF},
               aligns={'B': AL, 'O': AL, **{c: AR for c in 'EFGHIJKLM'}})
    for r in range(MN_R0, MN_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    _help(ws, MN_HDR, [(MN_IDX, '送货单第几行'), (MN_KEY, '品名规范写法'), (MN_LAST, '最后一次键'), (MN_MAX, '最高单价')], MN_R0, MN_R1)
    _help(ws, DN_R0 - 1, [(MS_HP, '单价（对送货单行）'), (MS_HK, '日期键（对送货单行）')], DN_R0, DN_R1)
    ws[f'{MN_IDX}{MN_HDR}'].font = F_HELP
    _neg_red(ws, f'E{MN_R0}:I{MN_R1}', f'E{MN_R0}')
    ws.freeze_panes = 'C4'
    print_setup(ws, None, landscape=True)
    ws.print_area = f'A1:P{MN_R1}'                # 打印到品名清单最后一行（不带右边隐藏的帮手列）


# ═══════════════════════════ 晚到单据（供应商送货单 / 外发加工单晚到） ═══════════════════════════
LT_HDR = 5
LT_R0 = 6                                     # 第 6 行＝上年送的（SM=0），第 7～18 行＝1～12 月送的
LT_RTOT = LT_R0 + 13
LT_SM = 'V'                                   # 隐藏：这一行的送货月 SM
LB_SEC = LT_RTOT + 2                          # ② 按供应商
LB_HDR = LB_SEC + 1
LB_TOT = LB_HDR + 1
LB_R0 = LB_TOT + 1
LB_N = UN_R1 - UN_R0 + 1
LB_R1 = LB_R0 + LB_N - 1
LB_IDX = 'V'                                  # 隐藏：第 k 家在【往来单位】第几行
LC_SEC, LC_HDR = LB_SEC, LB_HDR               # ③ 晚到清单（送货单）在右边 L～T 列
LC_R0, LC_N = LB_HDR + 1, 100
LC_R1 = LC_R0 + LC_N - 1
LD_SEC = LC_R1 + 2                            # ③ 晚到清单（外发）
LD_HDR = LD_SEC + 1
LD_R0, LD_N = LD_HDR + 1, 50
LD_R1 = LD_R0 + LD_N - 1
LC_IDX = 'W'                                  # 隐藏：第 k 笔在源表第几行


def build_late(wb, ctx):
    ws = wb[SH_LATE]
    widths(ws, {'A': 18, **{CL(2 + i): 10.5 for i in range(13)}, 'O': 12, 'P': 16, 'Q': 11, 'R': 16, 'S': 11, 'T': 12, 'U': 2})
    title(ws, '晚 到 单 据（供应商送货单、外发加工单晚到了多久 · 哪家最拖）', 'T', C_VIEW,
          '💡 全自动，看「送货单 / 外发加工单」拿到手比送货晚了多久。① 每一行是送货那个月，每一列是收到单子那个月：对角线上是当月就收到的，'
          '越往右越晚到；右边算出当月就收到、晚 1 个月、晚 2 个月以上的金额和晚到占比。'
          '「成本月份」规则：送货单按送货日期算到那个月（晚到的单录进去以后，那个月的成本会变）；'
          '【基础资料】⑩「每月结账日期」：x 月填了结账那天，x 月送的货、收单日期（空着按送货日期）晚于结账日期的，算到收单那个月，x 月的数不再变。'
          '最右边一列就是这种被挪了月份的金额（还有建账前送货、建账后才收到单的）。'
          '② 按供应商 / 加工厂：一共多少行、多少钱、平均晚几天、最长晚几天、晚到（≥提醒天数）的有多少。'
          '③ 晚到清单：滞后天数 ≥【基础资料】「收单滞后提醒天数」的单子，最晚的排最上面。建议每月月底催供应商把单子送来，再填结账日期。', h2=80)
    _home(ws, 'T')
    _lbl(ws, 'A3', '晚到提醒天数')
    _val(ws, 'B3', f'=N({LAGD})', '0" 天"')
    _lbl(ws, 'C3', '已结账到\n（基础资料⑩）', 'C3:D3')
    _val(ws, 'E3', f'=N({LOCK})', '0" 月";-0;"没结账"')       # 只是显示：最后一个填了结账日期的月份
    _lbl(ws, 'F3', '晚到的\n送货单', 'F3:G3')
    _val(ws, 'H3', f'=COUNT({dnr(DN_LATEK)})', '0" 行"')
    _lbl(ws, 'I3', '晚到的\n外发单', 'I3:J3')
    _val(ws, 'K3', f'=COUNT({otr(OT_LATEK)})', '0" 行"')
    _lbl(ws, 'L3', '晚到金额', 'L3:M3')
    _val(ws, 'N3', (f'=SUMIFS({dnr(DN_AMT)},{dnr(DN_OK)},1,{dnr(DN_LAG)},">="&N({LAGD}))'
                    f'+SUMIFS({otr(OT_AMT)},{otr(OT_OK)},1,{otr(OT_LAG)},">="&N({LAGD}))'), MONEY, 'N3:O3')
    _lbl(ws, 'P3', '晚 1 个月以上\n的金额占比')
    _val(ws, 'Q3', f'=S{LT_RTOT}', PCT)
    ws.row_dimensions[3].height = 32

    # ── ① 送货月 × 收单月
    section(ws, LT_HDR - 1, 'A', 'T', '① 送货月 × 收单月（元；送货单＋外发加工单，能记账的行）', H_DARK)
    put(ws, f'A{LT_HDR}', '送货月 ＼ 收单月', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, LT_HDR, 'B')
    put(ws, f'N{LT_HDR}', 13, F_HDR, fill(H_VIEW), '"次年"', ACW)          # 数字 13，显示「次年」
    for c, t in (('O', '合计'), ('P', '当月就收到'), ('Q', '晚 1 个月'), ('R', '晚 2 个月\n及以上'), ('S', '晚到占比\n（晚1个月以上）'),
                 ('T', '没算在送货月\n（结账 / 建账前）')):
        put(ws, f'{c}{LT_HDR}', t, F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[LT_HDR].height = 34
    MCOL = [CL(2 + i) for i in range(13)]       # B..N ＝ 收单 1～12 月、次年
    for sm in range(13):
        r = LT_R0 + sm
        V = f'${LT_SM}{r}'
        ws[f'{LT_SM}{r}'] = sm
        put(ws, f'A{r}', '上年送的（建账前）' if sm == 0 else f'{sm}月送的', F_TXTB, align=AL)
        for c in MCOL:
            ws[f'{c}{r}'] = (f'=SUMIFS({dnr(DN_AMT)},{dnr(DN_SM)},{V},{dnr(DN_RM)},{c}${LT_HDR},{dnr(DN_OK)},1)'
                             f'+SUMIFS({otr(OT_AMT)},{otr(OT_SM)},{V},{otr(OT_RM)},{c}${LT_HDR},{otr(OT_OK)},1)')
        ws[f'O{r}'] = f'=SUM(B{r}:N{r})'
        ws[f'P{r}'] = f'=IF({V}=0,0,INDEX($B{r}:$N{r},{V}))'
        ws[f'Q{r}'] = f'=INDEX($B{r}:$N{r},{V}+1)'
        ws[f'R{r}'] = f'=ROUND(O{r}-P{r}-Q{r},2)'
        ws[f'S{r}'] = f'=IF(O{r}=0,"",ROUND((Q{r}+R{r})/O{r},4))'
        ws[f'T{r}'] = (f'=SUMIFS({dnr(DN_AMT)},{dnr(DN_SM)},{V},{dnr(DN_OK)},1,{dnr(DN_CM)},"<>"&{V})'
                       f'+SUMIFS({otr(OT_AMT)},{otr(OT_SM)},{V},{otr(OT_OK)},1,{otr(OT_CM)},"<>"&{V})')
    cols = MCOL + list('OPQRST')
    style_rows(ws, LT_R0, LT_R0 + 12, cols, auto=cols, fmts={**{c: MONEY for c in cols}, 'S': PCT},
               aligns={c: AR for c in cols if c != 'S'})
    for r in range(LT_R0, LT_R0 + 13):
        ws[f'A{r}'].border = BD
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'O{r}'].fill = FILL_SUB
    # 对角线（当月就收到）浅绿，越往右越晚：晚 2 个月以上浅红
    rg = f'B{LT_R0}:N{LT_R0 + 12}'
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(B${LT_HDR}=${LT_SM}{LT_R0},B{LT_R0}<>0)'], fill=FILL_OK))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(B${LT_HDR}=${LT_SM}{LT_R0}+1,B{LT_R0}<>0)'], fill=FILL_YEL))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(B${LT_HDR}>=${LT_SM}{LT_R0}+2,B{LT_R0}<>0)'], fill=fill('FFFCE4D6')))
    r = LT_RTOT
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in MCOL + list('OPQRT'):
        put(ws, f'{c}{r}', f'=SUM({c}{LT_R0}:{c}{LT_R0 + 12})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'S{r}', f'=IF(O{r}=0,"",ROUND((Q{r}+R{r})/O{r},4))', F_AUTOB, FILL_TOT, PCT, AC)
    _help(ws, LT_HDR, [(LT_SM, '送货月')], LT_R0, LT_R0 + 12)

    # ── ② 按供应商 / 加工厂
    section(ws, LB_SEC, 'A', 'J', '② 按供应商 / 加工厂（能记账的行，全年）', H_DARK)
    heads = [('A', '供应商 / 加工厂'), ('B', '类型'), ('C', '行数'), ('D', '金额'), ('E', '平均晚几天\n（填了收单日期的）'),
             ('F', '最长\n晚几天'), ('G', '没填收单\n日期的行'), ('H', f'="≥"&N({LAGD})&"天"&CHAR(10)&"行数"'),
             ('I', f'="≥"&N({LAGD})&"天"&CHAR(10)&"金额"'), ('J', '晚到金额\n占比')]
    header(ws, LB_HDR, heads, H_VIEW, height=40)
    cnt = f'${LB_IDX}${LB_SEC}'
    ws[cnt.replace('$', '')] = f'=COUNT({unr(UN_APK)})'
    dn = lambda c: dnr(c)
    ot = lambda c: otr(c)
    for r in range(LB_R0, LB_R1 + 1):
        V = f'${LB_IDX}{r}'
        e = f'{V}=0'
        u = esc(f'$A{r}')
        k = f'ROW()-{LB_R0 - 1}'
        ws[f'{LB_IDX}{r}'] = f'=IF({k}>{cnt},0,SMALL({unr(UN_APK)},{k}))'

        lagv = {'DN': f',{dn(DN_LAGV)},">=0"', 'OT': f',{ot(OT_LAGV)},">=0"'}
        late = {'DN': f',{dn(DN_LAG)},">="&N({LAGD})', 'OT': f',{ot(OT_LAG)},">="&N({LAGD})'}
        cnt_all = f'COUNTIFS({dn(DN_UNITK)},{u},{dn(DN_OK)},1)+COUNTIFS({ot(OT_UNITK)},{u},{ot(OT_OK)},1)'
        amt_all = f'SUMIFS({dn(DN_AMT)},{dn(DN_UNITK)},{u},{dn(DN_OK)},1)+SUMIFS({ot(OT_AMT)},{ot(OT_UNITK)},{u},{ot(OT_OK)},1)'
        nd = (f'(COUNTIFS({dn(DN_UNITK)},{u},{dn(DN_OK)},1{lagv["DN"]})'
              f'+COUNTIFS({ot(OT_UNITK)},{u},{ot(OT_OK)},1{lagv["OT"]}))')
        sl = (f'(SUMIFS({dn(DN_LAG)},{dn(DN_UNITK)},{u},{dn(DN_OK)},1{lagv["DN"]})'
              f'+SUMIFS({ot(OT_LAG)},{ot(OT_UNITK)},{u},{ot(OT_OK)},1{lagv["OT"]}))')
        ws[f'A{r}'] = f'=IF({e},"",TRIM(INDEX({UN_NAMES},{V})&""))'
        ws[f'B{r}'] = f'=IF({e},"",INDEX({UN_TYPES_R},{V})&"")'
        ws[f'C{r}'] = f'=IF({e},"",{cnt_all})'
        ws[f'D{r}'] = f'=IF({e},"",{amt_all})'
        ws[f'E{r}'] = f'=IF({e},"",IF({nd}=0,"",ROUND({sl}/{nd},1)))'
        ws[f'F{r}'] = (f'=IF({e},"",MAX(SUMPRODUCT(MAX(({dn(DN_UNITK)}=$A{r})*({dn(DN_OK)}=1)*{dn(DN_LAG)})),'
                       f'SUMPRODUCT(MAX(({ot(OT_UNITK)}=$A{r})*({ot(OT_OK)}=1)*{ot(OT_LAG)}))))')
        ws[f'G{r}'] = f'=IF({e},"",C{r}-{nd})'
        ws[f'H{r}'] = (f'=IF({e},"",COUNTIFS({dn(DN_UNITK)},{u},{dn(DN_OK)},1{late["DN"]})'
                       f'+COUNTIFS({ot(OT_UNITK)},{u},{ot(OT_OK)},1{late["OT"]}))')
        ws[f'I{r}'] = (f'=IF({e},"",SUMIFS({dn(DN_AMT)},{dn(DN_UNITK)},{u},{dn(DN_OK)},1{late["DN"]})'
                       f'+SUMIFS({ot(OT_AMT)},{ot(OT_UNITK)},{u},{ot(OT_OK)},1{late["OT"]}))')
        ws[f'J{r}'] = f'=IF({e},"",IF(N(D{r})=0,"",ROUND(I{r}/D{r},4)))'
    cols = list('ABCDEFGHIJ')
    style_rows(ws, LB_R0, LB_R1, cols, auto=cols,
               fmts={'C': CNTF, 'D': MONEY, 'E': DAYF, 'F': CNTF, 'G': CNTF, 'H': CNTF, 'I': MONEY, 'J': PCT},
               aligns={'A': AL, 'D': AR, 'I': AR})
    for r in range(LB_R0, LB_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    r = LB_TOT
    put(ws, f'A{r}', f'="合计（"&{cnt}&" 家）"', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'B{r}', None, F_TXTB, FILL_TOT)
    for c, fm in (('C', CNTF), ('D', MONEY), ('G', CNTF), ('H', CNTF), ('I', MONEY)):
        put(ws, f'{c}{r}', f'=SUM({c}{LB_R0}:{c}{LB_R1})', F_AUTOB, FILL_TOT, fm, AR if fm == MONEY else AC)
    put(ws, f'E{r}', (f'=IF(COUNTIFS({dn(DN_OK)},1,{dn(DN_LAGV)},">=0")+COUNTIFS({ot(OT_OK)},1,{ot(OT_LAGV)},">=0")=0,"",'
                      f'ROUND((SUMIFS({dn(DN_LAG)},{dn(DN_OK)},1,{dn(DN_LAGV)},">=0")+SUMIFS({ot(OT_LAG)},{ot(OT_OK)},1,{ot(OT_LAGV)},">=0"))'
                      f'/(COUNTIFS({dn(DN_OK)},1,{dn(DN_LAGV)},">=0")+COUNTIFS({ot(OT_OK)},1,{ot(OT_LAGV)},">=0")),1))'),
        F_AUTOB, FILL_TOT, DAYF, AC)
    put(ws, f'F{r}', f'=IF(COUNT(F{LB_R0}:F{LB_R1})=0,"",MAX(F{LB_R0}:F{LB_R1}))', F_AUTOB, FILL_TOT, CNTF, AC)
    put(ws, f'J{r}', f'=IF(N(D{r})=0,"",ROUND(I{r}/D{r},4))', F_AUTOB, FILL_TOT, PCT, AC)
    ws.conditional_formatting.add(f'F{LB_R0}:F{LB_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER(F{LB_R0}),F{LB_R0}>=N({LAGD}))'], font=F_RED))
    _help(ws, LB_HDR, [(LB_IDX, '往来单位第几行')], LB_R0, LB_R1)
    ws[f'{LB_IDX}{LB_HDR}'].font = F_HELP

    # ── ③ 晚到清单（送货单 / 外发）
    def late_list(sec, hdr, r0, n, src, what):
        rngf, cDATE, cRDATE, cLAG, cUNIT, cNO, cNAME, cAMT, cCM, cLATEK = src
        section(ws, sec, 'L', 'T', f'③ 晚到清单：{what}（滞后 ≥ 提醒天数，最晚的在上，最多 {n} 行）', 'FFC00000')
        hd = [('L', '序号'), ('M', '送货日期' if what == '送货单' else '单据日期'), ('N', '收单日期'), ('O', '晚了\n几天'),
              ('P', '供应商' if what == '送货单' else '加工厂'), ('Q', '单号'), ('R', '品名' if what == '送货单' else '加工内容'),
              ('S', '金额'), ('T', '成本月份')]
        header(ws, hdr, hd, 'FFC00000')
        c_ = f'${LC_IDX}${hdr}'
        ws[c_.replace('$', '')] = f'=COUNT({rngf(cLATEK)})'
        ws[c_.replace('$', '')].font = F_HELP
        for r in range(r0, r0 + n):
            W = f'${LC_IDX}{r}'
            e = f'{W}=0'
            k = f'ROW()-{r0 - 1}'
            ws[f'{LC_IDX}{r}'] = f'=IF({k}>{c_},0,MOD(LARGE({rngf(cLATEK)},{k}),100000))'
            ix = lambda col: f'INDEX({rngf(col)},{W})'
            ws[f'L{r}'] = f'=IF({e},"",{k})'
            ws[f'M{r}'] = f'=IF({e},"",{ix(cDATE)})'
            ws[f'N{r}'] = f'=IF({e},"",{ix(cRDATE)})'
            ws[f'O{r}'] = f'=IF({e},"",{ix(cLAG)})'
            ws[f'P{r}'] = f'=IF({e},"",{ix(cUNIT)}&"")'
            ws[f'Q{r}'] = f'=IF({e},"",{ix(cNO)}&"")'
            ws[f'R{r}'] = f'=IF({e},"",{ix(cNAME)}&"")'
            ws[f'S{r}'] = f'=IF({e},"",{ix(cAMT)})'
            ws[f'T{r}'] = f'=IF({e},"",{ix(cCM)})'
            ws[f'{LC_IDX}{r}'].font = F_HELP
        cols_ = list('LMNOPQRST')
        style_rows(ws, r0, r0 + n - 1, cols_, auto=cols_,
                   fmts={'L': CNTF, 'M': DATE, 'N': DATE, 'O': CNTF, 'S': MONEY, 'T': MHDR}, aligns={'P': AL, 'R': AL, 'S': AR})
        for r in range(r0, r0 + n):
            for c in cols_:
                ws[f'{c}{r}'].fill = FILL_NONE
        put(ws, f'M{r0 + n}', f'=IF({c_}>{n},"⚠ 一共 "&{c_}&" 行，只列最晚的 {n} 行","")', F_RED, align=AL, border=False)
    late_list(LC_SEC, LC_HDR, LC_R0, LC_N, (dnr, DN_DATE, DN_RDATE, DN_LAG, DN_UNITK, DN_NO, DN_NAME, DN_AMT, DN_CM, DN_LATEK),
              '送货单')
    late_list(LD_SEC, LD_HDR, LD_R0, LD_N, (otr, OT_DATE, OT_RDATE, OT_LAG, OT_UNITK, OT_NO, OT_OP, OT_AMT, OT_CM, OT_LATEK),
              '外发')
    hide(ws, LC_IDX)
    ws.freeze_panes = 'B4'
    print_setup(ws, None, landscape=True)


# ═══════════════════════════ 应付账款汇总 ═══════════════════════════
AP_TOT, AP_HDR, AP_R0 = 4, 5, 6                # 合计行在表头上面（自动筛选只从表头第 5 行起，筛选时合计不会被藏掉）
AP_N = UN_R1 - UN_R0 + 1
AP_R1 = AP_R0 + AP_N - 1
AP_IDX, AP_LD, AP_LR, AP_LP = 'P', 'Q', 'R', 'S'          # 隐藏：往来单位第几行、最后送货 / 收单 / 付款日期（数字）
AP_MX = '$Q$3'                                            # 隐藏：实际用的月份（C3 空着＝最新月份）
AP_HD, AP_HR, AP_HOD, AP_HOR, AP_HP = 'U', 'V', 'W', 'X', 'Y'
# 隐藏帮手列（跟源表一行对一行，都是干净的数字，给 SUMPRODUCT(MAX()) 用）：
#   U/V ＝【送货单登记】能记账行的送货日期 / 收单日期；W/X ＝【外发加工登记】的；Y ＝【资金日记账】能记账、借方是 2202（付款）的日期
#   （这几列的第 5 行起就是源表的数据行，所以它们的小标题写在第 4 行）


def build_aps(wb, ctx):
    ws = wb[SH_APS]
    widths(ws, {'A': 5, 'B': 18, 'C': 10, 'D': 12, 'E': 13, 'F': 13, 'G': 13, 'H': 12, 'I': 12, 'J': 11, 'K': 11, 'L': 11,
                'M': 9, 'N': 40, 'O': 2})
    title(ws, '应 付 账 款 汇 总（每家供应商 / 加工厂：送了多少、付了多少、还欠多少）', 'N', C_VIEW,
          '💡 全自动，C3 选月份（黄格；空着＝最新月份）＝看到那个月月底为止。每家一行（【往来单位】里类型是材料供应商、外发加工厂的）：'
          '期初欠款（建账日）＋ 送货 / 加工（送货单、外发加工单记的应付，退货是负数）− 付款（资金日记账付的）＝ 月末还欠多少。'
          '欠款是负数＝付的比登记的送货单还多：多半是还有送货单没拿到（供应商隔几个月才拿单来），记得催单。'
          '金额都来自【记账分录】的应付账款（2202），所以手工分录调的账也算在里面。要跟某一家逐笔对账，去【供应商对账单】。'
          '第 5 行表头可以筛选（比如只看外发加工厂），第 4 行合计跟着筛选结果变；第 3 行的数一直是全部单位的。', h2=64)
    _home(ws, 'N')
    selector(ws, 'B3', '看到几月', 'C3', None, f'={AX_M}', MHDR, '看到哪个月月底（1～12）；空着＝最新月份')
    m = AP_MX
    ws[m.replace('$', '')] = f'={_mx("$C$3")}'
    ws[m.replace('$', '')].font = F_HELP
    mc = f'"<="&{m}'
    bal = (f'ROUND({COA_OPEN_OF("2202")}+{mat(chr(34) + "2202" + chr(34), "C", m, True)}'
           f'-{mat(chr(34) + "2202" + chr(34), "D", m, True)},2)')
    _lbl(ws, 'D3', '应付账款科目\n月末余额')
    _val(ws, 'E3', f'={bal}')
    _lbl(ws, 'F3', '下面清单\n合计（全部）')
    _val(ws, 'G3', f'=ROUND(SUM(G{AP_R0}:G{AP_R1}),2)')      # 全部单位（不跟筛选变），跟科目余额对账用
    _lbl(ws, 'H3', '不在清单里的\n（别的单位）')
    _val(ws, 'I3', f'=ROUND(E3-G3,2)')
    _lbl(ws, 'J3', '付多了的', 'J3:K3')
    _val(ws, 'L3', f'=COUNTIF(G{AP_R0}:G{AP_R1},"<0")', '0" 家"')
    put(ws, 'M3', (f'=IF({_mx_ok("$C$3")},"",IF(TRIM($C$3&"")="","C3 空着＝最新月份 "&{m}&" 月。",'
                   f'"⚠ C3 要选 1～12：现在按最新月份 "&{m}&" 月。"))'
                   f'&"「不在清单里的」＝应付账款记在别的类型单位（或没写单位）的，【明细账】选 2202 能看到"'), F_NOTE, KPI_FILL, align=ALW)
    ws.merge_cells('M3:N3')
    ws.row_dimensions[3].height = 44
    heads = [('A', '序号'), ('B', '供应商 / 加工厂'), ('C', '类型'), ('D', '期初欠款\n（建账日）'),
             ('E', f'="1～"&{m}&"月"&CHAR(10)&"送货 / 加工"'), ('F', f'="1～"&{m}&"月"&CHAR(10)&"付款"'),
             ('G', f'={m}&"月末"&CHAR(10)&"还欠"'), ('H', f'={m}&"月"&CHAR(10)&"送货 / 加工"'), ('I', f'={m}&"月"&CHAR(10)&"付款"'),
             ('J', '最后\n送货日期'), ('K', '最后\n收单日期'), ('L', '最后\n付款日期'), ('M', '平均晚几天\n收到单子'), ('N', '提示')]
    header(ws, AP_HDR, heads, H_VIEW, height=40)
    cnt = f'${AP_IDX}$3'
    ws[cnt.replace('$', '')] = f'=COUNT({unr(UN_APK)})'
    for src, rr, hd, hr, dcol, rcol, okcol in ((SH_DN, range(DN_R0, DN_R1 + 1), AP_HD, AP_HR, DN_DATE, DN_RDATE, DN_OK),
                                               (SH_OUT, range(OT_R0, OT_R1 + 1), AP_HOD, AP_HOR, OT_DATE, OT_RDATE, OT_OK)):
        for r in rr:
            x = lambda col: f'{q(src)}!${col}{r}'
            ws[f'{hd}{r}'] = f'=IF({x(okcol)}=1,INT({x(dcol)}),0)'
            ws[f'{hr}{r}'] = f'=IF({x(okcol)}=1,IF(ISNUMBER({x(rcol)}),INT({x(rcol)}),0),0)'
    for r in range(J_R0, J_R1 + 1):
        x = lambda col: f'{q(SH_CASH)}!${col}{r}'
        ws[f'{AP_HP}{r}'] = f'=IF(AND({x(J_OK)}=1,{x(J_DR)}="2202"),INT({x(J_DATE)}),0)'
    for r in range(AP_R0, AP_R1 + 1):
        Pc = f'${AP_IDX}{r}'
        e = f'{Pc}=0'
        B = f'$B{r}'
        u = esc(B)
        k = f'ROW()-{AP_R0 - 1}'
        ws[f'{AP_IDX}{r}'] = f'=IF({k}>{cnt},0,SMALL({unr(UN_APK)},{k}))'
        ws[f'A{r}'] = f'=IF({e},"",{k})'
        ws[f'B{r}'] = f'=IF({e},"",TRIM(INDEX({UN_NAMES},{Pc})&""))'
        ws[f'C{r}'] = f'=IF({e},"",INDEX({UN_TYPES_R},{Pc})&"")'
        ws[f'D{r}'] = f'=IF({e},"",N(INDEX({unr(UN_AP0)},{Pc})))'
        ws[f'E{r}'] = f'=IF({e},"",{je_sum(chr(34) + "2202" + chr(34), "C", mc, u)})'
        ws[f'F{r}'] = f'=IF({e},"",{je_sum(chr(34) + "2202" + chr(34), "D", mc, u)})'
        ws[f'G{r}'] = f'=IF({e},"",ROUND(D{r}+E{r}-F{r},2))'
        ws[f'H{r}'] = f'=IF({e},"",{je_sum(chr(34) + "2202" + chr(34), "C", m, u)})'
        ws[f'I{r}'] = f'=IF({e},"",{je_sum(chr(34) + "2202" + chr(34), "D", m, u)})'

        def last(urng, cmrng, hcol, r0, r1):
            return f'SUMPRODUCT(MAX(({urng}={B})*({cmrng}<={m})*${hcol}${r0}:${hcol}${r1}))'
        ws[f'{AP_LD}{r}'] = (f'=IF({e},0,MAX({last(dnr(DN_UNITK), dnr(DN_CM), AP_HD, DN_R0, DN_R1)},'
                             f'{last(otr(OT_UNITK), otr(OT_CM), AP_HOD, OT_R0, OT_R1)}))')
        ws[f'{AP_LR}{r}'] = (f'=IF({e},0,MAX({last(dnr(DN_UNITK), dnr(DN_CM), AP_HR, DN_R0, DN_R1)},'
                             f'{last(otr(OT_UNITK), otr(OT_CM), AP_HOR, OT_R0, OT_R1)}))')
        ws[f'{AP_LP}{r}'] = f'=IF({e},0,{last("TRIM(" + jr(J_UNIT) + chr(38) + chr(34) + chr(34) + ")", jr(J_CM), AP_HP, J_R0, J_R1)})'
        for c, h in (('J', AP_LD), ('K', AP_LR), ('L', AP_LP)):
            ws[f'{c}{r}'] = f'=IF(N(${h}{r})=0,"",INT(${h}{r}))'
        lagv = (f',{dnr(DN_LAGV)},">=0",{dnr(DN_CM)},{mc}', f',{otr(OT_LAGV)},">=0",{otr(OT_CM)},{mc}')
        nd = (f'(COUNTIFS({dnr(DN_UNITK)},{u},{dnr(DN_OK)},1{lagv[0]})+COUNTIFS({otr(OT_UNITK)},{u},{otr(OT_OK)},1{lagv[1]}))')
        sl = (f'(SUMIFS({dnr(DN_LAG)},{dnr(DN_UNITK)},{u},{dnr(DN_OK)},1{lagv[0]})'
              f'+SUMIFS({otr(OT_LAG)},{otr(OT_UNITK)},{u},{otr(OT_OK)},1{lagv[1]}))')
        ws[f'M{r}'] = f'=IF({e},"",IF({nd}=0,"",ROUND({sl}/{nd},1)))'
        ws[f'N{r}'] = (f'=IF({e},"",IF(G{r}<-0.005,"⚠ 付的比登记的送货单多 "&TEXT(-G{r},"#,##0.00")&" 元：可能还有送货单没拿到，催一下",'
                       f'IF(AND(ABS(G{r})<0.005,ABS(D{r})+ABS(E{r})>0),"√ 结清了","")))')
    cols = [CL(i) for i in range(1, 15)]
    style_rows(ws, AP_R0, AP_R1, cols, auto=cols,
               fmts={'A': CNTF, **{c: MONEY for c in 'DEFGHI'}, 'J': DATE, 'K': DATE, 'L': DATE, 'M': DAYF},
               aligns={'B': AL, 'N': AL, **{c: AR for c in 'DEFGHI'}}, bold=['G'])
    for r in range(AP_R0, AP_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'G{r}'].fill = FILL_SUB
    # 合计行（第 4 行，在表头上面）：用 SUBTOTAL，筛选以后只合计看得见的那几家
    r = AP_TOT
    vis = f'SUBTOTAL(102,A{AP_R0}:A{AP_R1})'          # 看得见的家数（A 列序号是数字）
    put(ws, f'A{r}', '', F_TXTB, FILL_TOT)
    put(ws, f'B{r}', f'=IF({vis}={cnt},"合计（"&{cnt}&" 家）","筛选出 "&{vis}&" 家（共 "&{cnt}&" 家）")', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'C{r}', '', F_TXTB, FILL_TOT)
    for c in 'DEFGHI':
        put(ws, f'{c}{r}', f'=ROUND(SUBTOTAL(109,{c}{AP_R0}:{c}{AP_R1}),2)', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'JKL':
        put(ws, f'{c}{r}', f'=IF(SUBTOTAL(102,{c}{AP_R0}:{c}{AP_R1})=0,"",SUBTOTAL(104,{c}{AP_R0}:{c}{AP_R1}))', F_AUTOB, FILL_TOT, DATE, AC)
    put(ws, f'M{r}', '', F_AUTOB, FILL_TOT)
    put(ws, f'N{r}', f'=IF(COUNTIF(N{AP_R0}:N{AP_R1},"⚠*")=0,"","⚠ 全部单位里 "&COUNTIF(N{AP_R0}:N{AP_R1},"⚠*")&" 家付多了")',
        F_AUTOB, FILL_TOT, align=AL)
    ws.row_dimensions[r].height = 22
    _neg_red(ws, f'D{AP_TOT}:I{AP_TOT}', f'D{AP_TOT}')
    _neg_red(ws, f'D{AP_R0}:I{AP_R1}', f'D{AP_R0}')
    _chk_cf(ws, f'N{AP_R0}:N{AP_R1}', f'$N{AP_R0}')
    _help(ws, AP_HDR, [(AP_IDX, '往来单位第几行'), (AP_LD, '最后送货'), (AP_LR, '最后收单'), (AP_LP, '最后付款')], AP_R0, AP_R1)
    # 跟源表一行对一行的帮手列：第 5 行起是数据，小标题写在第 4 行
    _help(ws, AP_TOT, [(AP_HD, '送货日期'), (AP_HR, '收单日期')], DN_R0, DN_R1)
    _help(ws, AP_TOT, [(AP_HOD, '外发单据日期'), (AP_HOR, '外发收单日期')], OT_R0, OT_R1)
    _help(ws, AP_TOT, [(AP_HP, '付款日期')], J_R0, J_R1)
    ws[f'{AP_IDX}{AP_HDR}'].font = F_HELP
    ws.auto_filter.ref = f'A{AP_HDR}:N{AP_R1}'
    ws.freeze_panes = f'C{AP_R0}'
    print_setup(ws, f'{AP_TOT}:{AP_HDR}', landscape=True)
    ws.print_area = f'A1:N{AP_R1}'


# ═══════════════════════════ 供应商对账单 ═══════════════════════════
SS_CAP = 300
SS_HDR, SS_OPEN = 7, 8
SS_R0 = SS_OPEN + 1
SS_R1 = SS_R0 + SS_CAP + 5 - 1               # 多 5 行给浮动的「合计、签字栏」
SS_SLOT, SS_SRC, SS_SI, SS_KIND = 'P', 'Q', 'R', 'S'      # 隐藏：记账分录槽位、来源、源表第几行（1 起）、行的种类


def _first_unit(ctx, tp):
    for row in ctx.get(SH_UNIT, []):
        if row.get(UN_TYPE) == tp:
            return row.get(UN_NAME)
    return None


def _float_cf(ws, rg, first_kind_cell):
    """浮动清单：1＝数据行（画框）、2＝合计行（画框、加粗、底色）、3/4＝签字栏（加粗）"""
    k = first_kind_cell
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{k}=1'], border=BD_CF))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{k}=2'], border=BD_CF, fill=FILL_TOT, font=F_BOLD))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'{k}>=3'], font=F_BOLD))


def build_sst(wb, ctx):
    ws = wb[SH_SST]
    widths(ws, {'A': 5, 'B': 11, 'C': 7, 'D': 11, 'E': 26, 'F': 8, 'G': 8, 'H': 6, 'I': 10, 'J': 13, 'K': 13, 'L': 13,
                'M': 11, 'N': 30, 'O': 2})
    title(ws, '供 应 商 对 账 单', 'N', C_VIEW,
          '💡 C3 选供应商 / 加工厂，F3、H3 填起止日期（黄格）。上面是汇总：期初欠款（起始日以前，含建账时的期初）＋ 本期送货 / 加工 − 本期付款 ＝ 期末欠款；'
          '下面按日期一笔一笔列出来（送货、退货、加工、付款、对方退款、手工调整），最后一列滚动余额，最下面是合计和签字栏。'
          '日期＝记账日期（一般就是送货单上的日期；那个月已结账、单子晚到被挪到收单月的，备注里写了原来的送货日期）。'
          '打印：A4 横向，打印区域已经设好（到清单最后一行），每页重复表头。最多列 300 笔，多了把日期范围缩短。', h2=48)
    _home(ws, 'N')
    sup = _first_unit(ctx, '材料供应商')
    selector(ws, 'B3', '供应商', 'C3', sup, f'={UN_NAMES}', '@', '选【往来单位】里的供应商 / 加工厂')
    ws.merge_cells('C3:D3')
    selector(ws, 'E3', '起始日期', 'F3', _dt.date(YEAR0, 1, 1), None, DATE, None)
    selector(ws, 'G3', '截止日期', 'H3', _dt.date(YEAR0, 12, 31), None, DATE, None)
    dv_date(ws, 'F3')
    dv_date(ws, 'H3')
    assert (SS_UNIT, SS_D1, SS_D2) == tuple(cell(SH_SST, x) for x in ('C3', 'F3', 'H3'))
    U = 'TRIM($C$3&"")'
    u = esc(U)
    d1 = 'N($F$3)'
    cnt_raw, cnt = f'${SS_SLOT}$5', f'${SS_SLOT}$6'
    ws[f'{SS_SLOT}5'] = f'=COUNT({jer(JE_SSK)})'
    ws[f'{SS_SLOT}6'] = f'=MIN({cnt_raw},{SS_CAP})'
    put(ws, 'I3', (f'=IF({U}="","← 先在 C3 选供应商",IF(COUNTIF({UN_NAMES},{u})=0,"⚠ 【往来单位】里没有这个单位",'
                   f'IF({cnt_raw}>{SS_CAP},"⚠ 这段时间有 "&{cnt_raw}&" 笔，只列前 {SS_CAP} 笔：把日期范围缩短",'
                   f'IF(NOT(ISNUMBER($H$3)),"⚠ 截止日期空着＝不限",""))))'), F_RED, align=AL, border=False)
    ws.merge_cells('I3:N3')
    # 第 4 行：单位信息
    urow = f'MATCH({u},{UN_NAMES},0)'
    info = lambda col: f'=IFERROR(INDEX({unr(col)},{urow})&"","")'
    _lbl(ws, 'A4', '我方', 'A4:B4')
    _val(ws, 'C4', f'={CO}&""', None, 'C4:E4', F_KPI_L, FILL_NONE)
    _lbl(ws, 'F4', '联系人')
    _val(ws, 'G4', info(UN_CONT), None, 'G4:H4', F_KPI_L, FILL_NONE)
    _lbl(ws, 'I4', '电话')
    _val(ws, 'J4', info(UN_TEL), '@', 'J4:K4', F_KPI_L, FILL_NONE)
    _lbl(ws, 'L4', '结算方式')
    _val(ws, 'M4', info(UN_TERM), None, 'M4:N4', F_KPI_L, FILL_NONE)
    # 第 5 行：汇总
    open_ = (f'ROUND(SUMIF({UN_NAMES},{u},{unr(UN_AP0)})'
             f'+SUMIFS({jer(JE_AMT)},{jer(JE_UNIT)},{u},{jer(JE_CR)},"2202",{jer(JE_DATE)},"<"&{d1})'
             f'-SUMIFS({jer(JE_AMT)},{jer(JE_UNIT)},{u},{jer(JE_DR)},"2202",{jer(JE_DATE)},"<"&{d1}),2)')
    _lbl(ws, 'A5', '期初欠款', 'A5:B5')
    _val(ws, 'C5', f'=IF({U}="",0,{open_})', MONEY, 'C5:D5')
    _lbl(ws, 'E5', '本期送货 / 加工')
    _val(ws, 'F5', f'=SUMIFS({jer(JE_AMT)},{jer(JE_SSK)},">0",{jer(JE_CR)},"2202")', MONEY, 'F5:G5')
    _lbl(ws, 'H5', '本期付款', 'H5:I5')
    _val(ws, 'J5', f'=SUMIFS({jer(JE_AMT)},{jer(JE_SSK)},">0",{jer(JE_DR)},"2202")', MONEY)
    _lbl(ws, 'K5', '期末欠款')
    _val(ws, 'L5', '=ROUND(C5+F5-J5,2)', MONEY, None, F_KPI_V)
    put(ws, 'M5', '期末＝期初＋送货−付款\n（正数＝我方欠贵方）', F_NOTE, KPI_FILL, align=ACW)
    ws.merge_cells('M5:N5')
    ws.row_dimensions[4].height = 22
    ws.row_dimensions[5].height = 32
    d2t = f'IF(ISNUMBER($H$3),{_dtxt_cn("$H$3")},"今天")'
    put(ws, 'A6', (f'=IF({U}="","",IF(L5>0.005,"截至 "&{d2t}&"，我方尚欠贵方 "&TEXT(L5,"#,##0.00")&" 元。",'
                   f'IF(L5<-0.005,"截至 "&{d2t}&"，我方多付贵方 "&TEXT(-L5,"#,##0.00")&" 元（贵方欠我方）。",'
                   f'"截至 "&{d2t}&"，双方结清。"))&"请核对下面明细，无误请签字盖章回传；有出入请注明。")'),
        F_TXTB, align=AL, border=False)
    ws.merge_cells('A6:N6')
    ws.row_dimensions[6].height = 22
    heads = [('A', '序号'), ('B', '日期'), ('C', '类型'), ('D', '单号'), ('E', '品名 / 内容'), ('F', '颜色'), ('G', '数量'),
             ('H', '单位'), ('I', '单价'), ('J', '送货 / 加工\n金额'), ('K', '付款金额'), ('L', '余额\n（我方欠）'), ('M', '收单日期'),
             ('N', '备注')]
    header(ws, SS_HDR, heads, H_VIEW)
    # 期初行
    r = SS_OPEN
    for c in 'ABCDEFGHIJKMN':
        put(ws, f'{c}{r}', None, F_TXTB, FILL_SUB)
    put(ws, f'B{r}', '=IF(ISNUMBER($F$3),$F$3,"")', F_TXTB, FILL_SUB, DATE, AC)
    put(ws, f'E{r}', f'="期初欠款"&IF(ISNUMBER($F$3),"（"&{_dtxt("$F$3")}&" 以前）","")', F_TXTB, FILL_SUB, align=AL)
    put(ws, f'L{r}', '=C5', F_AUTOB, FILL_SUB, MONEY, AR)

    src_base = (f'IF($Q{{r}}="日记账",{J_R0},IF($Q{{r}}="送货单",{DN_R0},IF($Q{{r}}="外发",{OT_R0},IF($Q{{r}}="手工",{MJ_R0},0))))')
    for r in range(SS_R0, SS_R1 + 1):
        Pc, Q, R, S = (f'${c}{r}' for c in (SS_SLOT, SS_SRC, SS_SI, SS_KIND))
        k = f'ROW()-{SS_R0 - 1}'
        je = lambda col: f'INDEX({jer(col)},{Pc})'
        dn = lambda col: f'INDEX({dnr(col)},{R})'
        ot = lambda col: f'INDEX({otr(col)},{R})'
        jn = lambda col: f'INDEX({jr(col)},{R})'
        mj = lambda col: f'INDEX({mjr(col)},{R})'
        isdn, isot = f'{Q}="送货单"', f'{Q}="外发"'
        d1_ = f'{S}<>1'
        f = {}
        f[SS_SLOT] = f'=IF({k}>{cnt},0,MOD(SMALL({jer(JE_SSK)},{k}),100000))'
        f[SS_SRC] = f'=IF({Pc}=0,"",{je(JE_SRC)}&"")'
        f[SS_SI] = f'=IF({Pc}=0,0,IF(N({je(JE_SROW)})=0,0,N({je(JE_SROW)})-{src_base.format(r=r)}+1))'
        f[SS_KIND] = f'=IF({Pc}>0,1,IF({k}={cnt}+1,2,IF({k}={cnt}+3,3,IF({k}={cnt}+5,4,0))))'
        f['A'] = f'=IF({S}=1,{k},"")'
        f['B'] = f'=IF({S}=1,{je(JE_DATE)},"")'
        f['C'] = (f'=IF({d1_},"",IF({isdn},IF(N({je(JE_AMT)})<0,"退货","送货"),IF({isot},"加工",'
                  f'IF({Q}="日记账",IF({je(JE_DR)}="2202","付款","对方退款"),"调整"))))')
        f['D'] = (f'=IF({d1_},"",IF({isdn},{dn(DN_NO)}&"",IF({isot},{ot(OT_NO)}&"",IF({Q}="日记账",{jn(J_NO)}&"",""))))')
        sty = f'TRIM({ot(OT_STY)}&"")'
        f['E'] = (f'=IF({S}=2,"本期合计",IF({S}=3,"供应商确认（签字盖章）：",IF({S}=4,"日期：　　　年　　月　　日",IF({d1_},"",IF({isdn},{dn(DN_NAME)}&"",IF({isot},TRIM({ot(OT_OP)}&"")'
                  f'&IF({sty}="","","（"&{sty}&"）"),{je(JE_MEMO)}&""))))))')
        f['F'] = f'=IF(AND({S}=1,{isdn}),{dn(DN_COLOR)}&"","")'
        nz = lambda x: f'IF(ISNUMBER({x}),{x},"")'
        f['G'] = f'=IF({d1_},"",IF({isdn},{nz(dn(DN_QTY))},IF({isot},{nz(ot(OT_QTY))},"")))'
        f['H'] = f'=IF({d1_},"",IF({isdn},{dn(DN_UNIT)}&"",IF({isot},{ot(OT_UNIT)}&"","")))'
        f['I'] = f'=IF({d1_},"",IF({isdn},{nz(dn(DN_PRICE))},IF({isot},{nz(ot(OT_PRICE))},"")))'
        f['J'] = f'=IF({S}=2,$F$5,IF({d1_},"",IF({je(JE_CR)}="2202",{je(JE_AMT)},"")))'
        f['K'] = f'=IF({S}=2,$J$5,IF({d1_},"",IF({je(JE_DR)}="2202",{je(JE_AMT)},"")))'
        f['L'] = f'=IF({S}=2,$L$5,IF({d1_},"",ROUND(N(L{r - 1})+N(J{r})-N(K{r}),2)))'
        rd = f'IF({isdn},{dn(DN_RDATE)},{ot(OT_RDATE)})'
        f['M'] = f'=IF({d1_},"",IF(OR({isdn},{isot}),IF(ISNUMBER({rd}),{rd},""),""))'
        sd = f'IF({isdn},{dn(DN_DATE)},{ot(OT_DATE)})'
        note = (f'IF({isdn},{dn(DN_NOTE)}&"",IF({isot},{ot(OT_NOTE)}&"",IF({Q}="日记账",{jn(J_NOTE)}&"",'
                f'IF({Q}="手工",{mj(MJ_NOTE)}&"",""))))')
        f['N'] = (f'=IF({S}=3,"本厂确认（签字）：",IF({S}=4,"日期：　　　年　　月　　日",IF({d1_},"",IF(OR({isdn},{isot}),IF(INT({sd})<>$B{r},IF({isdn},"送货日期 ","单据日期 ")&{_dtxt(sd)}&"；",""),"")'
                  f'&{note})))')
        for col, v in f.items():
            ws[f'{col}{r}'] = v
    cols = [CL(i) for i in range(1, 15)]
    _plain_rows(ws, SS_R0, SS_R1, cols,
                fmts={'A': CNTF, 'B': DATE, 'G': QTYF, 'I': MONEY, 'J': MONEY, 'K': MONEY, 'L': MONEY, 'M': DATE},
                aligns={'B': AL, 'E': AL, 'N': AL, 'G': AR, 'I': AR, 'J': AR, 'K': AR, 'L': AR})
    _float_cf(ws, f'A{SS_R0}:N{SS_R1}', f'${SS_KIND}{SS_R0}')
    _neg_red(ws, f'J{SS_R0}:L{SS_R1}', f'J{SS_R0}')
    _help(ws, SS_HDR, [(SS_SLOT, '分录槽位'), (SS_SRC, '来源'), (SS_SI, '源表第几行'), (SS_KIND, '行种类')], SS_R0, SS_R1)
    ws[f'{SS_SLOT}5'].font = ws[f'{SS_SLOT}6'].font = F_HELP
    ws.freeze_panes = f'A{SS_HDR + 1}'
    print_setup(ws, f'{SS_HDR}:{SS_HDR}', landscape=True)
    ws.print_area = f'A1:N{SS_R1}'                # 到清单容量末行（含浮动的合计、签字栏），不带右边隐藏列


# ═══════════════════════════ 客户对账单 ═══════════════════════════
CS_HDR = 8
CS_R0 = CS_HDR + 1
CS_LN, CS_RN, CS_ON = 400, 100, 100           # 交货明细 / 收款等 / 按订单号小计 最多几行
CS_L1 = CS_R0 + CS_LN                         # 多 1 行给浮动「合计」
CS_RR1 = CS_R0 + CS_RN
CS_O1 = CS_R0 + CS_ON
CS_LI, CS_LK, CS_RS, CS_RK, CS_OF, CS_OI, CS_OK_, CS_ONO, CS_OH = 'U', 'V', 'W', 'X', 'Y', 'Z', 'AA', 'AB', 'AC'
# 隐藏：U 交货第 k 行在订单明细第几行  V 种类 ｜ W 收款等第 k 笔的分录槽位  X 种类 ｜
#       Y（跟【订单明细】一行对一行）这一行是不是这张单在清单里第一次出现（是＝本行序号 i）
#       AC（跟【订单明细】一行对一行）在清单里的行＝这张单的 OFR（单在订单明细第一次出现的序号，按 NOK 规范写法认单），否则 ""
#       Z 第 k 张单在清单里第一次出现是订单明细第几行  AA 种类  AB 这张单的 OFR（按它分组小计）
#   「第一次出现」只在同一张单的行里找：从这张单在订单明细的首行（OFR）找到本行，用 MATCH（找到第一个就停），不从表头数起


def build_cst(wb, ctx):
    ws = wb[SH_CST]
    widths(ws, {'A': 5, 'B': 11, 'C': 10, 'D': 10, 'E': 14, 'F': 7, 'G': 9, 'H': 12, 'I': 2,
                'J': 5, 'K': 11, 'L': 26, 'M': 12, 'N': 12, 'O': 2, 'P': 11, 'Q': 7, 'R': 9, 'S': 12, 'T': 2})
    title(ws, '客 户 对 账 单（给电商部对账）', 'S', C_VIEW,
          '💡 C3 选客户，F3、H3 填起止日期（黄格）。上面是汇总：期初应收（起始日以前没收回的，含建账时的期初）＋ 本期交货 − 本期收款 ＋ 其他（退款给客户等）'
          '＝ 期末应收（正数＝客户还欠我们）。左边是这段时间每一笔交货（【订单明细】填了交货日期的行，退回来的是负数），右边是收款等'
          '（资金日记账收的货款、退的钱、手工分录调的），最右边按订单号（采购单号）小计，方便跟电商部一张单一张单对。'
          '金额都是净数（退货已经扣掉、返修重交照样算钱）；「交货双数」不含退货（负数）和返修重交的行（备注写了「返修」的），'
          '跟【订单汇总】【款式成本利润】的已交双数一个口径，所以左边清单「数量」逐行加起来可能跟它不一样。'
          '打印：A4 横向，打印区域已经设好，每页重复表头。', h2=64)
    _home(ws, 'S')
    cus = _first_unit(ctx, '客户')
    selector(ws, 'B3', '客户', 'C3', cus, f'={UN_NAMES}', '@', '选【往来单位】里的客户')
    ws.merge_cells('C3:D3')
    selector(ws, 'E3', '起始日期', 'F3', _dt.date(YEAR0, 1, 1), None, DATE, None)
    selector(ws, 'G3', '截止日期', 'H3', _dt.date(YEAR0, 12, 31), None, DATE, None)
    dv_date(ws, 'F3')
    dv_date(ws, 'H3')
    assert (CS_UNIT, CS_D1, CS_D2) == tuple(cell(SH_CST, x) for x in ('C3', 'F3', 'H3'))
    U = 'TRIM($C$3&"")'
    u = esc(U)
    d1 = 'N($F$3)'
    lc_raw, lc = f'${CS_LI}$6', f'${CS_LI}$7'
    rc_raw, rc = f'${CS_RS}$6', f'${CS_RS}$7'
    oc_raw, oc = f'${CS_OI}$6', f'${CS_OI}$7'
    ws[f'{CS_LI}6'], ws[f'{CS_LI}7'] = f'=COUNT({odr(OD_CSK)})', f'=MIN({lc_raw},{CS_LN})'
    ws[f'{CS_RS}6'], ws[f'{CS_RS}7'] = f'=COUNT({jer(JE_CSK)})', f'=MIN({rc_raw},{CS_RN})'
    ws[f'{CS_OI}6'] = f'=COUNT(${CS_OF}${OD_R0}:${CS_OF}${OD_R1})'
    ws[f'{CS_OI}7'] = f'=MIN({oc_raw},{CS_ON})'
    put(ws, 'I3', (f'=IF({U}="","← 先在 C3 选客户",IF(COUNTIF({UN_NAMES},{u})=0,"⚠ 【往来单位】里没有这个客户",'
                   f'IF(OR(NOT(ISNUMBER($F$3)),NOT(ISNUMBER($H$3))),"⚠ 起止日期都要填（没填的话交货明细列不出来）",'
                   f'IF(OR({lc_raw}>{CS_LN},{rc_raw}>{CS_RN},{oc_raw}>{CS_ON}),"⚠ 笔数太多，只列了前面的：把日期范围缩短",""))))'),
        F_RED, align=AL, border=False)
    ws.merge_cells('I3:S3')
    # 第 4、5 行：汇总
    open_ = (f'ROUND(SUMIF({UN_NAMES},{u},{unr(UN_AR0)})'
             f'+SUMIFS({odr(OD_AMT)},{odr(OD_CUSTK)},{u},{odr(OD_OK)},1,{odr(OD_DDATE)},"<"&{d1})'
             f'+SUMIFS({jer(JE_AMT)},{jer(JE_UNIT)},{u},{jer(JE_DR)},"1122",{jer(JE_SRC)},"<>收入",{jer(JE_DATE)},"<"&{d1})'
             f'-SUMIFS({jer(JE_AMT)},{jer(JE_UNIT)},{u},{jer(JE_CR)},"1122",{jer(JE_SRC)},"<>收入",{jer(JE_DATE)},"<"&{d1}),2)')
    for lc_, lab, vc, f, merge_l, merge_v, font in (
            ('B4', '期初应收\n（起始日以前）', 'B5', f'=IF({U}="",0,{open_})', 'B4:C4', 'B5:C5', F_VAL),
            ('D4', '本期交货', 'D5', f'=SUMIFS({odr(OD_AMT)},{odr(OD_CSK)},">0")', 'D4:E4', 'D5:E5', F_VAL),
            ('F4', '本期收款', 'F5', f'=SUMIFS({jer(JE_AMT)},{jer(JE_CSK)},">0",{jer(JE_CR)},"1122")', 'F4:G4', 'F5:G5', F_VAL),
            ('H4', '其他\n（退款给客户等）', 'H5', f'=SUMIFS({jer(JE_AMT)},{jer(JE_CSK)},">0",{jer(JE_DR)},"1122")', None, None, F_VAL),
            ('K4', '期末应收', 'K5', '=ROUND(B5+D5-F5+H5,2)', 'K4:L4', 'K5:L5', F_KPI_V)):
        _lbl(ws, lc_, lab, merge_l)
        _val(ws, vc, f, MONEY, merge_v, font)
    put(ws, 'M4', '期末＝期初＋本期交货−本期收款＋其他\n（正数＝客户还欠我们）', F_NOTE, KPI_FILL, align=ACW)
    ws.merge_cells('M4:N5')
    put(ws, 'P4', '本期交货双数\n不含退货、返修', F_KPI_L, KPI_FILL, align=ACW)
    ws.merge_cells('P4:Q4')
    _val(ws, 'P5', f'=SUMIFS({odr(OD_CQ)},{odr(OD_CSK)},">0")', CNTF, 'P5:Q5')      # 成本双数 CQ：跟订单汇总、款式成本利润一个口径
    put(ws, 'R4', '本期交货\n订单数', F_KPI_L, KPI_FILL, align=ACW)
    ws.merge_cells('R4:S4')
    _val(ws, 'R5', f'={oc_raw}', CNTF, 'R5:S5')
    ws.row_dimensions[4].height = 30
    ws.row_dimensions[5].height = 24
    d2t = f'IF(ISNUMBER($H$3),{_dtxt_cn("$H$3")},"今天")'
    put(ws, 'A6', (f'=IF({U}="","",IF(K5>0.005,"截至 "&{d2t}&"，"&{U}&" 尚欠货款 "&TEXT(K5,"#,##0.00")&" 元。",'
                   f'IF(K5<-0.005,"截至 "&{d2t}&"，"&{U}&" 多付了 "&TEXT(-K5,"#,##0.00")&" 元（预收）。",'
                   f'"截至 "&{d2t}&"，双方结清。")))'), F_TXTB, align=AL, border=False)
    ws.merge_cells('A6:S6')
    section(ws, 7, 'A', 'H', '交货明细（【订单明细】交了货的行，按交货日期）', H_DARK)
    section(ws, 7, 'J', 'N', '收款等（应收账款的收付、调整）', H_DARK)
    section(ws, 7, 'P', 'S', '按订单号小计', H_DARK)
    header(ws, CS_HDR, [('A', '序号'), ('B', '交货日期'), ('C', '订单号'), ('D', '款式'), ('E', '颜色及规格'), ('F', '数量'),
                        ('G', '单价'), ('H', '金额'), ('J', '序号'), ('K', '日期'), ('L', '摘要'), ('M', '收款'), ('N', '其他\n（退款等）'),
                        ('P', '订单号'), ('Q', '交货\n行数'), ('R', '交货双数\n不含退货\n和返修'), ('S', '金额')], H_VIEW, height=46)
    # 左：交货明细（合计行的「数量」＝清单逐行加起来的净数，含退货负数、返修重交）
    dq_net = f'SUMIFS({odr(OD_DQ)},{odr(OD_CSK)},">0")'
    for r in range(CS_R0, CS_L1 + 1):
        Uc, V = f'${CS_LI}{r}', f'${CS_LK}{r}'
        k = f'ROW()-{CS_R0 - 1}'
        od = lambda col: f'INDEX({odr(col)},{Uc})'
        d = f'{V}<>1'
        ws[f'{CS_LI}{r}'] = f'=IF({k}>{lc},0,MOD(SMALL({odr(OD_CSK)},{k}),100000))'
        ws[f'{CS_LK}{r}'] = f'=IF({Uc}>0,1,IF({k}={lc}+1,2,0))'
        ws[f'A{r}'] = f'=IF({d},"",{k})'
        ws[f'B{r}'] = f'=IF({V}=2,"合计",IF({d},"",INT({od(OD_DDATE)})))'
        ws[f'C{r}'] = f'=IF({d},"",{od(OD_NOK)}&"")'
        ws[f'D{r}'] = f'=IF({d},"",{od(OD_STY)}&"")'
        ws[f'E{r}'] = f'=IF({d},"",{od(OD_SPEC)}&"")'
        ws[f'F{r}'] = f'=IF({V}=2,{dq_net},IF({d},"",{od(OD_DQ)}))'
        ws[f'G{r}'] = f'=IF({d},"",{od(OD_PRICEU)})'
        ws[f'H{r}'] = f'=IF({V}=2,$D$5,IF({d},"",{od(OD_AMT)}))'
    # 右：收款等
    for r in range(CS_R0, CS_RR1 + 1):
        W, X = f'${CS_RS}{r}', f'${CS_RK}{r}'
        k = f'ROW()-{CS_R0 - 1}'
        je = lambda col: f'INDEX({jer(col)},{W})'
        d = f'{X}<>1'
        ws[f'{CS_RS}{r}'] = f'=IF({k}>{rc},0,MOD(SMALL({jer(JE_CSK)},{k}),100000))'
        ws[f'{CS_RK}{r}'] = f'=IF({W}>0,1,IF({k}={rc}+1,2,0))'
        ws[f'J{r}'] = f'=IF({d},"",{k})'
        ws[f'K{r}'] = f'=IF({X}=2,"合计",IF({d},"",{je(JE_DATE)}))'
        ws[f'L{r}'] = f'=IF({d},"",{je(JE_MEMO)}&"")'
        ws[f'M{r}'] = f'=IF({X}=2,$F$5,IF({d},"",IF({je(JE_CR)}="1122",{je(JE_AMT)},"")))'
        ws[f'N{r}'] = f'=IF({X}=2,$H$5,IF({d},"",IF({je(JE_DR)}="1122",{je(JE_AMT)},"")))'
    # 这张单第一次在清单里出现（跟【订单明细】一行对一行；只在这一行在清单里时才查，别的行直接空）
    #   AC＝在清单里的行放这张单的 OFR（数字），Y＝从这张单的首行（OFR）往下 MATCH 第一个在清单里的同单行，就是本行 → 本行序号 i
    oh = f'${CS_OH}${OD_R0}:${CS_OH}${OD_R1}'
    for r in range(OD_R0, OD_R1 + 1):
        csk = f'{q(SH_ORD)}!${OD_CSK}{r}'
        ofr = f'{q(SH_ORD)}!${OD_OFR}{r}'
        o = f'${CS_OH}{r}'
        ws[f'{CS_OH}{r}'] = f'=IF({csk}="","",N({ofr}))'
        ws[f'{CS_OF}{r}'] = (f'=IF({o}="","",IF({o}<1,ROW()-{OD_HDR},'
                             f'IF(MATCH({o},INDEX({oh},{o}):${CS_OH}${OD_R1},0)=ROW()-{OD_HDR}-{o}+1,ROW()-{OD_HDR},"")))')
    # 右右：按订单号小计（按 OFR 分组＝按订单号规范写法 NOK 认单：数字 / 文字 / 前后空格都算同一张）
    for r in range(CS_R0, CS_O1 + 1):
        Z, AA, AB = f'${CS_OI}{r}', f'${CS_OK_}{r}', f'${CS_ONO}{r}'
        k = f'ROW()-{CS_R0 - 1}'
        d = f'{AA}<>1'
        crit = f'{odr(OD_OFR)},{AB},{odr(OD_CSK)},">0"'
        ws[f'{CS_OI}{r}'] = f'=IF({k}>{oc},0,SMALL(${CS_OF}${OD_R0}:${CS_OF}${OD_R1},{k}))'
        ws[f'{CS_OK_}{r}'] = f'=IF({Z}>0,1,IF({k}={oc}+1,2,0))'
        ws[f'{CS_ONO}{r}'] = f'=IF({Z}=0,0,N(INDEX({odr(OD_OFR)},{Z})))'
        ws[f'P{r}'] = f'=IF({AA}=2,"合计",IF({d},"",INDEX({odr(OD_NOK)},{Z})&""))'
        ws[f'Q{r}'] = f'=IF({AA}=2,{lc_raw},IF({d},"",COUNTIFS({crit})))'
        ws[f'R{r}'] = f'=IF({AA}=2,$P$5,IF({d},"",SUMIFS({odr(OD_CQ)},{crit})))'
        ws[f'S{r}'] = f'=IF({AA}=2,$D$5,IF({d},"",SUMIFS({odr(OD_AMT)},{crit})))'
    _plain_rows(ws, CS_R0, CS_L1, list('ABCDEFGH'), fmts={'A': CNTF, 'B': DATE, 'F': QTYF, 'G': MONEY, 'H': MONEY},
                aligns={'B': AL, 'E': AL, 'F': AR, 'G': AR, 'H': AR})
    _plain_rows(ws, CS_R0, CS_RR1, list('JKLMN'), fmts={'J': CNTF, 'K': DATE, 'M': MONEY, 'N': MONEY},
                aligns={'K': AL, 'L': AL, 'M': AR, 'N': AR})
    _plain_rows(ws, CS_R0, CS_O1, list('PQRS'), fmts={'Q': CNTF, 'R': CNTF, 'S': MONEY}, aligns={'P': AL, 'S': AR})
    _float_cf(ws, f'A{CS_R0}:H{CS_L1}', f'${CS_LK}{CS_R0}')
    _float_cf(ws, f'J{CS_R0}:N{CS_RR1}', f'${CS_RK}{CS_R0}')
    _float_cf(ws, f'P{CS_R0}:S{CS_O1}', f'${CS_OK_}{CS_R0}')
    _neg_red(ws, f'F{CS_R0}:H{CS_L1}', f'F{CS_R0}')
    _help(ws, CS_HDR, [(CS_LI, '订单明细第几行'), (CS_LK, '种类'), (CS_RS, '分录槽位'), (CS_RK, '种类'),
                       (CS_OI, '订单首行'), (CS_OK_, '种类'), (CS_ONO, '这张单的 OFR')], CS_R0, CS_L1)
    _help(ws, OD_HDR, [(CS_OF, '本单在清单里首次（对订单明细行）'), (CS_OH, '清单里的行：本单 OFR')], OD_R0, OD_R1)
    for c in (CS_LI, CS_RS, CS_OI):
        ws[f'{c}6'].font = ws[f'{c}7'].font = F_HELP
    ws.freeze_panes = f'A{CS_R0}'
    print_setup(ws, f'{CS_HDR}:{CS_HDR}', landscape=True)
    ws.print_area = f'A1:S{max(CS_L1, CS_RR1, CS_O1)}'      # 到三块清单容量最长那块的末行，不带右边隐藏列


# ═══════════════════════════ 费用汇总 ═══════════════════════════
FE_HDR_ = 4                                   # 表头（月份是数字）；第 5～64 行 ↔【基础资料】④ 第 5～64 行
assert FE_R0 == FE_HDR_ + 1
FX_WAGE = FE_R1 + 1                           # 工资 × 制造 / 管理 / 销售
FX_DEP = FX_WAGE + 3                          # 折旧费 × 制造 / 管理 / 销售
FX_NONE = FX_DEP + 3                          # 未分项目 × 4101 / 5601 / 5602 / 5603
FX_LAST = FX_NONE + 3
FX_SEC = FX_LAST + 2
FX_C0 = FX_SEC + 1                            # 7 个归类小计
FX_PER = FX_C0 + len(FE_CLASSES)              # 期间费用合计
FX_ALL = FX_PER + 1                           # 费用合计（7 类）
FX_CHK = FX_ALL + 1                           # 核对：制造费用 vs 成本分摊表
FX_CODE = 'Q'                                 # 隐藏：科目编码


def build_fee(wb, ctx):
    ws = wb[SH_FEE]
    widths(ws, {'A': 16, 'B': 11, **{CL(3 + i): 10.5 for i in range(12)}, 'O': 13, 'P': 2})
    title(ws, '费 用 汇 总（每个费用项目每个月花了多少）', 'O', C_VIEW,
          '💡 全自动。每个费用项目（【基础资料】④）一行，1～12 月，来自【记账分录】（资金日记账的「费用支出」、手工分录填了费用项目的）。'
          '工资、折旧按「制造 / 管理 / 销售」单列；没填费用项目（或填的项目跟科目对不上）的算「未分项目」。下面按归类小计。'
          '制造费用（车间房租、水电、车间管理人员工资、车间折旧……）会按交货双数分到各款式的成本里；'
          '管理、销售、财务费用不分到款式，直接进【利润表】（【款式成本利润】里按收入比例摊的「期间费用」就是它们）。退回来的钱是负数。', h2=64)
    _home(ws, 'O')
    put(ws, f'A{FE_HDR_}', '费用项目', F_HDR, fill(H_VIEW), align=ACW)
    put(ws, f'B{FE_HDR_}', '归类', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, FE_HDR_, 'C')
    put(ws, f'O{FE_HDR_}', '全年', F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[FE_HDR_].height = 30
    ws[f'{FX_CODE}{FE_HDR_}'] = '科目'
    MC = [CL(3 + i) for i in range(12)]

    def amt(r, c):
        a = f'$A{r}'
        crit = f'{jer(JE_FEE)},{esc(f"TRIM({a})")},'
        return (f'=IF(OR({a}="",${FX_CODE}{r}=""),"",ROUND(SUMIFS({jer(JE_AMT)},{crit}{jer(JE_DR)},${FX_CODE}{r},{jer(JE_M)},{c}${FE_HDR_})'
                f'-SUMIFS({jer(JE_AMT)},{crit}{jer(JE_CR)},${FX_CODE}{r},{jer(JE_M)},{c}${FE_HDR_}),2))')
    for r in range(FE_R0, FE_R1 + 1):
        ws[f'A{r}'] = f"={q(SH_BASE)}!${FE_NAME}${r}&\"\""
        ws[f'B{r}'] = f"={q(SH_BASE)}!${FE_CLS}${r}&\"\""
        ws[f'{FX_CODE}{r}'] = f"={q(SH_BASE)}!${FE_CODE}${r}&\"\""
    cls_code = {k: v[0] for k, v in FE_CLASSES.items()}
    code_cls = {v: k for k, v in cls_code.items()}
    specials = ([(FX_WAGE + i, '工资', k) for i, k in enumerate(('制造费用', '管理费用', '销售费用'))]
                + [(FX_DEP + i, '折旧费', k) for i, k in enumerate(('制造费用', '管理费用', '销售费用'))])
    for r, nm, k in specials:
        ws[f'A{r}'], ws[f'B{r}'], ws[f'{FX_CODE}{r}'] = nm, k, cls_code[k]
    for r in range(FE_R0, FX_DEP + 3):
        for c in MC:
            ws[f'{c}{r}'] = amt(r, c)
        ws[f'O{r}'] = f'=IF(OR($A{r}="",${FX_CODE}{r}=""),"",SUM(C{r}:N{r}))'
    # 未分项目：科目（含下级）本月借−贷（科目×月矩阵 mat()；制造费用不算月末转出到生产成本的那笔）− 上面各项目里这个科目的
    #   制造转入那笔＝【记账分录】制造转入块第 m 个槽位的金额（借 400104 贷 4101，layout JE_BLK）
    xfer0 = JE_BLK['制造转入'][0]
    for i, code in enumerate((MOH, '5601', '5602', '5603')):
        r = FX_NONE + i
        ws[f'A{r}'], ws[f'B{r}'], ws[f'{FX_CODE}{r}'] = '未分项目', code_cls[code], code
        Qc = f'${FX_CODE}{r}'
        for m, c in enumerate(MC, 1):
            back = f'+N({q(SH_JE)}!${JE_AMT}${xfer0 + m - 1})' if code == MOH else ''
            ws[f'{c}{r}'] = (f'=ROUND({mat(Qc + chr(38) + chr(34) + "*" + chr(34), "D", m)}'
                             f'-{mat(Qc + chr(38) + chr(34) + "*" + chr(34), "C", m)}{back}'
                             f'-SUMIFS({c}${FE_R0}:{c}${FX_DEP + 2},${FX_CODE}${FE_R0}:${FX_CODE}${FX_DEP + 2},{Qc}),2)')
        ws[f'O{r}'] = f'=SUM(C{r}:N{r})'
    cols = ['A', 'B'] + MC + ['O']
    style_rows(ws, FE_R0, FX_LAST, cols, auto=cols, fmts={c: MONEY for c in MC + ['O']},
               aligns={'A': AL, 'B': AC, **{c: AR for c in MC + ['O']}})
    for r in range(FE_R0, FX_LAST + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'O{r}'].fill = FILL_SUB
        ws[f'O{r}'].font = F_AUTOB
    for r in range(FX_WAGE, FX_LAST + 1):
        for c in ('A', 'B'):
            ws[f'{c}{r}'].fill = fill('FFF2F2F2')
    # 按归类小计
    section(ws, FX_SEC, 'A', 'O', '按归类小计', H_DARK)
    for i, (k, (code, _cf)) in enumerate(FE_CLASSES.items()):
        r = FX_C0 + i
        put(ws, f'A{r}', k, F_TXTB, FILL_SUB, align=AL)
        put(ws, f'B{r}', code, F_TXT, FILL_SUB, '@', AC)
        for c in MC:
            put(ws, f'{c}{r}', f'=SUMIF($B${FE_R0}:$B${FX_LAST},$A{r},{c}${FE_R0}:{c}${FX_LAST})', F_AUTO, FILL_SUB, MONEY, AR)
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})', F_AUTOB, FILL_SUB, MONEY, AR)
    rows_of = {k: FX_C0 + i for i, k in enumerate(FE_CLASSES)}
    r = FX_PER
    put(ws, f'A{r}', '期间费用合计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'B{r}', '管理＋销售＋财务', F_NOTE, FILL_TOT, align=AC)
    for c in MC + ['O']:
        put(ws, f'{c}{r}', f'={c}{rows_of["管理费用"]}+{c}{rows_of["销售费用"]}+{c}{rows_of["财务费用"]}', F_AUTOB, FILL_TOT, MONEY, AR)
    r = FX_ALL
    put(ws, f'A{r}', '费用合计（7 类）', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'B{r}', None, F_NOTE, FILL_TOT)
    for c in MC + ['O']:
        put(ws, f'{c}{r}', f'=SUM({c}{FX_C0}:{c}{FX_PER - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    r = FX_CHK
    put(ws, f'A{r}', '核对：制造费用', F_TXTB, align=AL)
    put(ws, f'B{r}', '−成本分摊表', F_NOTE, align=AC)
    for i, c in enumerate(MC):
        put(ws, f'{c}{r}', f"=ROUND({c}{rows_of['制造费用']}-N({al('制造', '本月发生', i + 1)}),2)", F_AUTOB, None, CHKF, AC)
    put(ws, f'O{r}', f'=ROUND(SUM(C{r}:N{r}),2)', F_AUTOB, None, CHKF, AC)
    _chk_num_cf(ws, f'C{r}:O{r}', f'C{r}')
    put(ws, f'A{r + 1}', '核对：制造费用小计应该等于【成本分摊表】「制造 本月发生」（制造费用全部分到款式成本里）。'
                         '「未分项目」不是 0 的：去【记账分录】按科目筛选看是哪几笔没填费用项目。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{r + 1}:O{r + 1}')
    ws.row_dimensions[r + 1].height = 30
    # 第 3 行
    _lbl(ws, 'A3', '本年制造费用')
    _val(ws, 'B3', f'=O{rows_of["制造费用"]}', MONEY, 'B3:C3')
    _lbl(ws, 'D3', '本年期间费用\n（管理＋销售＋财务）', 'D3:E3')
    _val(ws, 'F3', f'=O{FX_PER}', MONEY, 'F3:G3')
    _lbl(ws, 'H3', '其中\n未分项目', 'H3:I3')
    _val(ws, 'J3', f'=SUM(O{FX_NONE}:O{FX_LAST})', MONEY, 'J3:K3')
    put(ws, 'L3', f'=IF(ROUND(O{FX_CHK},2)=0,"√ 制造费用跟成本分摊表对上了","✗ 制造费用跟成本分摊表差 "&TEXT(O{FX_CHK},"#,##0.00"))',
        F_KPI_L, KPI_FILL, align=ACW)
    ws.merge_cells('L3:O3')
    _chk_cf(ws, 'L3', '$L$3')
    ws.row_dimensions[3].height = 32
    _neg_red(ws, f'C{FE_R0}:O{FX_ALL}', f'C{FE_R0}')
    ws.conditional_formatting.add(f'A{FX_NONE}:O{FX_LAST}', FormulaRule(
        formula=[f'ROUND($O{FX_NONE},2)<>0'], fill=FILL_YEL))
    _help(ws, FE_HDR_, [(FX_CODE, '科目')], FE_R0, FX_LAST)
    ws.freeze_panes = f'C{FE_R0}'
    print_setup(ws, f'{FE_HDR_}:{FE_HDR_}', landscape=True)


# ═══════════════════════════ 工资汇总 ═══════════════════════════
WS_HDR = 4                                    # 第 5～24 行 ↔【基础资料】⑤ 第 5～24 行
assert DP_R0 == WS_HDR + 1
WS_OTH = DP_R1 + 1                            # 部门对不上的
WS_TOT = WS_OTH + 1                           # 应发合计
WS_PAID = WS_TOT + 1                          # 实发
WS_OWE = WS_PAID + 1                          # 月末未发
WS_SEC = WS_OWE + 2
WS_D0 = WS_SEC + 1                            # 去向 4 行
WS_DT = WS_D0 + len(DP_TYPES)                 # 去向合计
WS_CODE = 'Q'


def build_wsum(wb, ctx):
    ws = wb[SH_WSUM]
    widths(ws, {'A': 16, 'B': 11, **{CL(3 + i): 10.5 for i in range(12)}, 'O': 13, 'P': 2})
    title(ws, '工 资 汇 总（各部门每月应发 · 实发 · 还欠多少工资）', 'O', C_VIEW,
          '💡 全自动。上面每个部门（【基础资料】⑤）每个月的应发工资（【工资登记】，按「月份」那一列算，不管哪天发的）；'
          '下面：应发合计、实发（资金日记账记的「付工资」，按付钱那天的月份）、月末还没发的工资（应付职工薪酬余额＝年初没发的＋累计应发−累计实发）。'
          '最下面按去向分：直接人工、制造费用（这两块分到款式成本里）、管理费用、销售费用。', h2=48)
    _home(ws, 'O')
    put(ws, f'A{WS_HDR}', '部门', F_HDR, fill(H_VIEW), align=ACW)
    put(ws, f'B{WS_HDR}', '类型', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, WS_HDR, 'C')
    put(ws, f'O{WS_HDR}', '全年', F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[WS_HDR].height = 30
    MC = [CL(3 + i) for i in range(12)]
    W2211 = '"2211*"'                             # 应付职工薪酬（含下级）：科目×月矩阵 mat() 取累计借贷
    for r in range(DP_R0, DP_R1 + 1):
        a = f'$A{r}'
        ws[f'A{r}'] = f"={q(SH_BASE)}!${DP_NAME}${r}&\"\""
        ws[f'B{r}'] = f"={q(SH_BASE)}!${DP_TYPE}${r}&\"\""
        for c in MC:
            ws[f'{c}{r}'] = f'=IF({a}="","",SUMIFS({wgr(WG_AMT)},{wgr(WG_DEPT)},{a},{wgr(WG_CM)},{c}${WS_HDR},{wgr(WG_OK)},1))'
        ws[f'O{r}'] = f'=IF({a}="","",SUM(C{r}:N{r}))'
    cols = ['A', 'B'] + MC + ['O']
    style_rows(ws, DP_R0, DP_R1, cols, auto=cols, fmts={c: MONEY for c in MC + ['O']},
               aligns={'A': AL, **{c: AR for c in MC + ['O']}})
    for r in range(DP_R0, DP_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'O{r}'].fill = FILL_SUB
    rows = [
        (WS_OTH, '其他（部门写法跟⑤对不上的）', '', lambda c: f'=ROUND({c}{WS_TOT}-SUM({c}{DP_R0}:{c}{DP_R1}),2)', None, 'sum'),
        (WS_TOT, '应发合计', '工资登记', lambda c: f'=SUMIFS({wgr(WG_AMT)},{wgr(WG_CM)},{c}${WS_HDR},{wgr(WG_OK)},1)', FILL_TOT, 'sum'),
        (WS_PAID, '实发（付工资）', '资金日记账',
         lambda c: f'=-SUMIFS({jr(J_NET)},{jr(J_CATX)},"付工资",{jr(J_CM)},{c}${WS_HDR},{jr(J_OK)},1)', FILL_SUB, 'sum'),
        (WS_OWE, '月末还没发的工资', '应付职工薪酬余额',
         lambda c: (f'=ROUND({COA_OPEN_OF("2211")}+{mat(W2211, "C", MC.index(c) + 1, True)}'
                    f'-{mat(W2211, "D", MC.index(c) + 1, True)},2)'),
         FILL_SUB, 'last'),
    ]
    for r, lab, note, fm, fl, yr in rows:
        put(ws, f'A{r}', lab, F_TXTB if fl else F_TXT, fl, align=AL)
        put(ws, f'B{r}', note, F_NOTE, fl, align=AC)
        for c in MC:
            put(ws, f'{c}{r}', fm(c), F_AUTOB if fl else F_AUTO, fl, MONEY, AR)
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})' if yr == 'sum' else f'=N{r}', F_AUTOB, fl or FILL_SUB, MONEY, AR)
    section(ws, WS_SEC, 'A', 'O', '按去向（工资进哪里）', H_DARK)
    for i, (tp, (_comp, code, dest)) in enumerate(DP_TYPES.items()):
        r = WS_D0 + i
        put(ws, f'A{r}', dest, F_TXTB, align=AL)
        put(ws, f'B{r}', code, F_NOTE, None, '@', AC)
        ws[f'{WS_CODE}{r}'] = code
        for c in MC:
            put(ws, f'{c}{r}', f'=SUMIFS({wgr(WG_AMT)},{wgr(WG_DRC)},${WS_CODE}{r},{wgr(WG_CM)},{c}${WS_HDR},{wgr(WG_OK)},1)',
                F_AUTO, None, MONEY, AR)
        put(ws, f'O{r}', f'=SUM(C{r}:N{r})', F_AUTOB, FILL_SUB, MONEY, AR)
    r = WS_DT
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'B{r}', '＝应发合计', F_NOTE, FILL_TOT, align=AC)
    for c in MC + ['O']:
        put(ws, f'{c}{r}', f'=SUM({c}{WS_D0}:{c}{WS_DT - 1})', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'A{r + 1}', '直接人工＝生产成本（计件按款式记的直接算到那个款）；制造费用＝车间辅助人员（分到款式）；管理、销售费用不算成本。'
                         '「其他」不是 0：【工资登记】里有部门名称跟【基础资料】⑤ 写法不一样（比如多了空格）的行。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{r + 1}:O{r + 1}')
    ws.row_dimensions[r + 1].height = 30
    _help(ws, WS_HDR, [(WS_CODE, '科目')], WS_D0, WS_DT - 1)
    # 第 3 行
    owe_now = f'ROUND({COA_OPEN_OF("2211")}+{mat(W2211, "C", 12, True)}-{mat(W2211, "D", 12, True)},2)'   # 全年累计＝现在
    _lbl(ws, 'A3', '年初没发的工资')
    _val(ws, 'B3', f'={COA_OPEN_OF("2211")}', MONEY, 'B3:C3')
    _lbl(ws, 'D3', '本年应发', 'D3:E3')
    _val(ws, 'F3', f'=O{WS_TOT}', MONEY, 'F3:G3')
    _lbl(ws, 'H3', '本年实发', 'H3:I3')
    _val(ws, 'J3', f'=O{WS_PAID}', MONEY, 'J3:K3')
    _lbl(ws, 'L3', '现在还欠工资', 'L3:M3')
    _val(ws, 'N3', f'={owe_now}', MONEY, 'N3:O3', F_KPI_V)
    ws.row_dimensions[3].height = 30
    _neg_red(ws, f'C{DP_R0}:O{WS_DT}', f'C{DP_R0}')
    ws.conditional_formatting.add(f'C{WS_OTH}:O{WS_OTH}', FormulaRule(formula=[f'ROUND(C{WS_OTH},2)<>0'], fill=FILL_YEL))
    ws.freeze_panes = f'C{DP_R0}'
    print_setup(ws, f'{WS_HDR}:{WS_HDR}', landscape=True)


# ═══════════════════════════ 账户余额表 ═══════════════════════════
AB_HDR = 4                                    # A 块：第 5～12 行 ↔【基础资料】② 第 5～12 行
assert AC_R0 == AB_HDR + 1
AB_TOT = AC_R1 + 1
NAC = AC_R1 - AC_R0 + 1
AB_BSEC = AB_TOT + 2                          # B 块：各账户月末余额
AB_BHDR = AB_BSEC + 1
AB_B0 = AB_BHDR + 1
AB_BTOT = AB_B0 + NAC
AB_CSEC = AB_BTOT + 2                         # C 块：收支类别 × 月
AB_CHDR = AB_CSEC + 1
AB_C0 = AB_CHDR + 1
NCT = CT_R1 - CT_R0 + 1
AB_CTOT = AB_C0 + NCT
AB_OM, AB_LM = 'P3', 'Q3'                     # 隐藏：建账月、最新月份（② 块只显示这两个月之间的）


def build_accb(wb, ctx):
    ws = wb[SH_ACCB]
    widths(ws, {'A': 14, 'B': 11, 'C': 11, 'D': 12, 'E': 12, 'F': 12, 'G': 12, 'H': 12, 'I': 12, 'J': 12, 'K': 12, 'L': 12,
                'M': 12, 'N': 40, 'O': 2})
    title(ws, '账 户 余 额 表（每个资金账户：收了多少、付了多少、每月月底还剩多少）', 'N', C_VIEW,
          '💡 全自动，来自【资金日记账】。① 每个账户：期初＋本年收入−本年支出＋内部转入−内部转出＝当前余额（跟资金日记账、【基础资料】② 一样）；'
          '「记账的余额」是会计科目里的数（只算校验不是 ✗ 的行），两个不一样就是日记账里有 ✗ 的行没进账（或手工分录动了资金科目），提示列会说。'
          '「实际余额」是在【基础资料】② 手填的银行 App / 支付宝上的数，对不上就是漏记、记错了。'
          '② 每个账户每个月月底的余额（含账户之间转账的两头；建账以前的月份、最新月份以后还没到的月份空着）。③ 每个收支类别每个月的净额（收进来是正数、付出去是负数；自己账户之间转账不算）。', h2=64)
    _home(ws, 'N')
    heads = [('A', '账户'), ('B', '类型'), ('C', '科目编码'), ('D', '期初余额\n（建账日）'), ('E', '本年收入\n（不含内部转入）'),
             ('F', '本年支出\n（不含内部转出）'), ('G', '内部转入'), ('H', '内部转出'), ('I', '当前余额\n（按日记账）'),
             ('J', '记账的余额\n（科目余额）'), ('K', '差额\n（当前−记账）'), ('L', '实际余额\n（手填）'), ('M', '实际−当前'), ('N', '提示')]
    header(ws, AB_HDR, heads, H_VIEW, height=40)
    base = lambda col, r: f"{q(SH_BASE)}!${col}${r}"
    net, acc, catx, tox = jr(J_NET), jr(J_ACC), jr(J_CATX), jr(J_TOX)
    for r in range(AC_R0, AC_R1 + 1):
        a = f'$A{r}'
        e = f'{a}=""'
        ws[f'A{r}'] = f'={base(AC_NAME, r)}&""'
        ws[f'B{r}'] = f'={base(AC_TYPE, r)}&""'
        ws[f'C{r}'] = f'=TRIM({base(AC_CODE, r)}&"")'
        ws[f'D{r}'] = f'=IF({e},"",N({base(AC_OPEN, r)}))'
        ws[f'E{r}'] = f'=IF({e},"",SUMIFS({net},{acc},{a},{catx},"<>{XFER}",{net},">0"))'
        ws[f'F{r}'] = f'=IF({e},"",-SUMIFS({net},{acc},{a},{catx},"<>{XFER}",{net},"<0"))'
        ws[f'G{r}'] = f'=IF({e},"",SUMIFS({net},{acc},{a},{catx},"{XFER}",{net},">0")-SUMIFS({net},{tox},{a},{net},"<0"))'
        ws[f'H{r}'] = f'=IF({e},"",-SUMIFS({net},{acc},{a},{catx},"{XFER}",{net},"<0")+SUMIFS({net},{tox},{a},{net},">0"))'
        ws[f'I{r}'] = f'=IF({e},"",N({base(AC_NOW, r)}))'
        ws[f'J{r}'] = (f'=IF(OR({e},$C{r}=""),"",ROUND(D{r}+{mat(f"$C{r}", "D", 12, True)}'
                       f'-{mat(f"$C{r}", "C", 12, True)},2))')          # 科目×月矩阵：全年累计
        ws[f'K{r}'] = f'=IF(J{r}="","",ROUND(I{r}-J{r},2))'
        ws[f'L{r}'] = f'=IF(OR({e},{base(AC_REAL, r)}=""),"",N({base(AC_REAL, r)}))'
        ws[f'M{r}'] = f'=IF(L{r}="","",ROUND(L{r}-I{r},2))'
        ws[f'N{r}'] = (f'=IF({e},"",IF($C{r}="","✗ 【基础资料】② 没选类型：算不出科目编码",'
                       f'IF(ROUND(ROUND(I{r},2)-ROUND(D{r}+E{r}-F{r}+G{r}-H{r},2),2)<>0,"✗ 收支加起来跟当前余额对不上（请告诉做表的人）",'
                       f'IF(ABS(N(K{r}))>=0.005,"⚠ 当前余额跟记账的差 "&TEXT(K{r},"#,##0.00")&"：日记账里有 ✗ 的行没进账（改好就对上了），或手工分录动了这个科目",'
                       f'IF(AND(L{r}<>"",ABS(N(M{r}))>=0.005),"⚠ 实际余额跟当前余额差 "&TEXT(M{r},"#,##0.00")&"：漏记或记错了，去资金日记账查","√")))))')
    cols = [CL(i) for i in range(1, 15)]
    style_rows(ws, AC_R0, AC_R1, cols, auto=cols, fmts={**{c: MONEY for c in 'DEFGHIJKLM'}, 'C': '@'},
               aligns={'A': AL, 'N': AL, **{c: AR for c in 'DEFGHIJKLM'}}, bold=['I'])
    for r in range(AC_R0, AC_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'I{r}'].fill = FILL_SUB
    r = AB_TOT
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in 'BC':
        put(ws, f'{c}{r}', None, F_TXTB, FILL_TOT)
    for c in 'DEFGHIJKLM':
        put(ws, f'{c}{r}', f'=ROUND(SUM({c}{AC_R0}:{c}{AC_R1}),2)', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'N{r}', f'=IF(COUNTIF(N{AC_R0}:N{AC_R1},"✗*")+COUNTIF(N{AC_R0}:N{AC_R1},"⚠*")=0,"√",'
                     f'"有提示的账户 "&(COUNTIF(N{AC_R0}:N{AC_R1},"✗*")+COUNTIF(N{AC_R0}:N{AC_R1},"⚠*"))&" 个")', F_AUTOB, FILL_TOT, align=AL)
    _chk_cf(ws, f'N{AC_R0}:N{AB_TOT}', f'$N{AC_R0}')
    _neg_red(ws, f'D{AC_R0}:M{AB_TOT}', f'D{AC_R0}')
    # 第 3 行
    _lbl(ws, 'A3', '总余额（当前）')
    _val(ws, 'B3', f'=I{AB_TOT}', MONEY, 'B3:C3', F_KPI_V)
    _lbl(ws, 'D3', '记账的\n资金余额')
    _val(ws, 'E3', f'=J{AB_TOT}', MONEY, 'E3:F3')
    _lbl(ws, 'G3', '本年收入\n（不含内部）')
    _val(ws, 'H3', f'=E{AB_TOT}', MONEY, 'H3:I3')
    _lbl(ws, 'J3', '本年支出\n（不含内部）')
    _val(ws, 'K3', f'=F{AB_TOT}', MONEY, 'K3:L3')
    put(ws, 'M3', f'=N{AB_TOT}', F_KPI_L, KPI_FILL, align=ACW)
    ws.merge_cells('M3:N3')
    _chk_cf(ws, 'M3', '$M$3')
    ws.row_dimensions[3].height = 32

    # ── ② 各账户月末余额：建账以前的月份、最新月份以后的月份显示空（不是 0：那些月份还没有数）
    put(ws, AB_OM, f'={OPEN_M}', F_HELP, border=False)            # 建账月
    put(ws, AB_LM, f'={LATEST}', F_HELP, border=False)            # 最新月份（有业务的最后一个月；还没业务＝建账月）
    hide(ws, AB_OM[0], AB_LM[0])
    om, lm = f'${AB_OM[0]}${AB_OM[1:]}', f'${AB_LM[0]}${AB_LM[1:]}'
    section(ws, AB_BSEC, 'A', 'N', '② 各账户每月月底的余额（含账户之间转账的两头；日期不对的行不算；建账以前、最新月份以后的月份不显示）', H_DARK)
    put(ws, f'A{AB_BHDR}', '账户', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, AB_BHDR, 'B')
    put(ws, f'N{AB_BHDR}', '当前余额', F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[AB_BHDR].height = 30
    MB = [CL(2 + i) for i in range(12)]
    off = AB_B0 - AC_R0
    for i in range(NAC):
        r = AB_B0 + i
        a = f'$A{r}'
        ws[f'A{r}'] = f'=A{r - off}'
        for c in MB:
            mm = f'{c}${AB_BHDR}'
            ws[f'{c}{r}'] = (f'=IF(OR({a}="",{mm}<{om},{mm}>{lm}),"",ROUND(N($D{r - off})+SUMIFS({net},{acc},{a},{jr(J_CM)},">=1",{jr(J_CM)},"<="&{mm})'
                             f'-SUMIFS({net},{tox},{a},{jr(J_CM)},">=1",{jr(J_CM)},"<="&{mm}),2))')
        ws[f'N{r}'] = f'=IF({a}="","",I{r - off})'
    cols = ['A'] + MB + ['N']
    style_rows(ws, AB_B0, AB_B0 + NAC - 1, cols, auto=cols, fmts={c: MONEY for c in MB + ['N']},
               aligns={'A': AL, **{c: AR for c in MB + ['N']}})
    for r in range(AB_B0, AB_B0 + NAC):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'N{r}'].fill = FILL_SUB
    r = AB_BTOT
    put(ws, f'A{r}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in MB + ['N']:
        put(ws, f'{c}{r}', f'=IF(COUNT({c}{AB_B0}:{c}{AB_BTOT - 1})=0,"",ROUND(SUM({c}{AB_B0}:{c}{AB_BTOT - 1}),2))',
            F_AUTOB, FILL_TOT, MONEY, AR)
    _neg_red(ws, f'B{AB_B0}:N{AB_BTOT}', f'B{AB_B0}')

    # ── ③ 收支类别 × 月
    section(ws, AB_CSEC, 'A', 'N', '③ 收支类别 × 月 净额（能记账的行；收进来正、付出去负；内部转账不算收支）', H_DARK)
    put(ws, f'A{AB_CHDR}', '收支类别', F_HDR, fill(H_VIEW), align=ACW)
    _mhdr(ws, AB_CHDR, 'B')
    put(ws, f'N{AB_CHDR}', '全年', F_HDR, fill(H_VIEW), align=ACW)
    ws.row_dimensions[AB_CHDR].height = 30
    offc = AB_C0 - CT_R0
    for i in range(NCT):
        r = AB_C0 + i
        a = f'$A{r}'
        ws[f'A{r}'] = f'={base(CT_NAME, r - offc)}&""'
        for c in MB:
            ws[f'{c}{r}'] = (f'=IF({a}="","",IF({a}="{XFER}",0,SUMIFS({net},{catx},{a},{jr(J_CM)},{c}${AB_CHDR},'
                             f'{jr(J_OK)},1)))')
        ws[f'N{r}'] = f'=IF({a}="","",SUM(B{r}:M{r}))'
    cols = ['A'] + MB + ['N']
    style_rows(ws, AB_C0, AB_CTOT - 1, cols, auto=cols, fmts={c: MONEY for c in MB + ['N']},
               aligns={'A': AL, **{c: AR for c in MB + ['N']}})
    for r in range(AB_C0, AB_CTOT):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
        ws[f'N{r}'].fill = FILL_SUB
    r = AB_CTOT
    put(ws, f'A{r}', '净额合计', F_TXTB, FILL_TOT, align=AC)
    for c in MB + ['N']:
        put(ws, f'{c}{r}', f'=ROUND(SUM({c}{AB_C0}:{c}{AB_CTOT - 1}),2)', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'A{r + 1}', '净额合计＝这个月所有账户加起来多了（正）或少了（负）多少钱；每个月：上月底总余额＋净额合计＝本月底总余额（只算能记账的行）。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{r + 1}:N{r + 1}')
    _neg_red(ws, f'B{AB_C0}:N{AB_CTOT}', f'B{AB_C0}')
    ws.freeze_panes = 'B4'
    print_setup(ws, None, landscape=True)


def build(wb, ctx):
    build_msum(wb, ctx)
    build_late(wb, ctx)
    build_aps(wb, ctx)
    build_sst(wb, ctx)
    build_cst(wb, ctx)
    build_fee(wb, ctx)
    build_wsum(wb, ctx)
    build_accb(wb, ctx)
