# -*- coding: utf-8 -*-
"""总账几张（第二阶段）：【记账分录】【科目余额表】【明细账】【利润表】【资产负债表】【现金流量表】。
   ★ 记账分录的列、隐藏接口列 O～S、各选择格（GL_* / BS_M / TB_M / BS_CHECK）的位置都按 layout.py 一字不差地实现。
   记账分录的来源块（日记账 / 送货单 / 外发 / 手工）第 i 个槽位直接引用来源表第 i 行（不用 INDEX），
   系统块（收入 / 工资 / 折旧 / 制造转入 / 成本结转）按 layout 注释的槽位排法，每月一组。"""
import re
from openpyxl.formatting.rule import FormulaRule
from common import *
from layout import *          # 注意：layout 的 C_RPT 等颜色覆盖 common 的同名常量

YR, OPEN, CO = P['YEAR'], P['OPEN'], P['CO']
KPI_FILL = fill('FFD9E1F2')
F_VAL = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_GOOD = Font(name=YH, sz=11, bold=True, color='FF00B050')
F_BAD = Font(name=YH, sz=11, bold=True, color='FFC00000')
F_HEADB = Font(name=YH, sz=10, bold=True, color='FF833C0C')
FILL_HEAD = fill('FFFBE5D6')
HDR = C_RPT                                   # 报表组表头颜色（深红棕）
MONTHF = '0"月";-0;""'


def _kpi(ws, coord_l, label, coord_v, formula, fmt=None, font=F_VAL):
    put(ws, coord_l, label, F_KPI_L, KPI_FILL, align=ACW)
    put(ws, coord_v, formula, font, KPI_FILL, fmt, AC)


def _good_bad(ws, coord):
    """√ 开头绿字、✗ 开头红底"""
    ws.conditional_formatting.add(coord, FormulaRule(formula=[f'LEFT({coord.split(":")[0]},1)="✗"'], fill=FILL_WARN, font=F_BAD))
    ws.conditional_formatting.add(coord, FormulaRule(formula=[f'LEFT({coord.split(":")[0]},1)="√"'], font=F_GOOD))


# ═══════════════════════════ 记账分录 ═══════════════════════════
SRC = {  # 来源块 → (来源表, 数据第一行)
    '日记账': (SH_CASH, J_R0), '送货单': (SH_DN, DN_R0), '外发': (SH_OUT, OT_R0), '手工': (SH_MJ, MJ_R0)}
WAGE_K = [('400103', '直接人工'), (MOH, '制造费用'), ('5602', '管理费用'), ('5601', '销售费用')]   # 工资块 k＝0～3
DEP_K = ['车间', '管理', '销售']                                                               # 折旧块 k＝0～2


def _src_row(name, sr, r):
    """来源块第 i 个槽位（来源表第 sr 行，本表第 r 行）的各列公式（不含 A、G/I 名称、O～S，那些每块一样）"""
    sh = SRC[name][0]
    x = lambda col: f"{q(sh)}!{col}{sr}"
    f = {}
    if name == '日记账':
        ok = f'{x(J_OK)}=1'
        f[JE_DATE] = f'=IF({ok},INT({x(J_DATE)}),"")'
        f[JE_M] = f'=IF({ok},{x(J_CM)},"")'
        memo = f'IF(TRIM({x(J_MEMO)}&"")="",{x(J_CATX)}&"",TRIM({x(J_MEMO)}&""))'
        f[JE_MEMO] = f'=IF({ok},{memo}&IF(TRIM({x(J_UNIT)}&"")="",""," · "&TRIM({x(J_UNIT)}&"")),"")'
        f[JE_DR] = f'=IF({ok},{x(J_DR)}&"","")'
        f[JE_CR] = f'=IF({ok},{x(J_CR)}&"","")'
        f[JE_AMT] = f'=IF({ok},N({x(J_AMT)}),0)'
        f[JE_UNIT] = f'=IF({ok},TRIM({x(J_UNIT)}&""),"")'
        f[JE_FEE] = f'=IF(AND({ok},{x(J_FCLS)}<>""),TRIM({x(J_FEE)}&""),"")'
        f[JE_STY] = f'=IF(AND({ok},LEFT({x(J_CKEY)},2)="款|"),TRIM({x(J_STY)}&""),"")'
    elif name == '送货单':
        ok = f'{x(DN_OK)}=1'
        f[JE_DATE] = f'=IF({ok},{x(DN_PDATE)},"")'
        f[JE_M] = f'=IF({ok},{x(DN_CM)},"")'
        nm, col, qty, un, no = (x(c) for c in (DN_NAME, DN_COLOR, DN_QTY, DN_UNIT, DN_NO))
        f[JE_MEMO] = (f'=IF({ok},"送货 "&TRIM({nm}&"")&IF(TRIM({col}&"")="",""," "&TRIM({col}&""))'
                      f'&IF(ISNUMBER({qty})," "&{qty}&TRIM({un}&""),"")'
                      f'&IF(TRIM({no}&"")="","","（单号 "&TRIM({no}&"")&"）"),"")')
        f[JE_DR] = f'=IF({ok},"{COMP_CODE["材料"]}","")'
        f[JE_CR] = f'=IF({ok},"2202","")'
        f[JE_AMT] = f'=IF({ok},N({x(DN_AMT)}),0)'
        f[JE_UNIT] = f'=IF({ok},{x(DN_UNITK)}&"","")'
        f[JE_STY] = f'=IF(AND({ok},{x(DN_ATYPE)}="款"),TRIM({x(DN_STY)}&""),"")'
    elif name == '外发':
        ok = f'{x(OT_OK)}=1'
        f[JE_DATE] = f'=IF({ok},{x(OT_PDATE)},"")'
        f[JE_M] = f'=IF({ok},{x(OT_CM)},"")'
        op, qty, un, sty = (x(c) for c in (OT_OP, OT_QTY, OT_UNIT, OT_STY))
        f[JE_MEMO] = (f'=IF({ok},"外发 "&TRIM({op}&"")&IF(ISNUMBER({qty})," "&{qty}&TRIM({un}&""),"")'
                      f'&IF(TRIM({sty}&"")="","","（"&TRIM({sty}&"")&"）"),"")')
        f[JE_DR] = f'=IF({ok},"{COMP_CODE["外发"]}","")'
        f[JE_CR] = f'=IF({ok},"2202","")'
        f[JE_AMT] = f'=IF({ok},N({x(OT_AMT)}),0)'
        f[JE_UNIT] = f'=IF({ok},{x(OT_UNITK)}&"","")'
        f[JE_STY] = f'=IF(AND({ok},{x(OT_ATYPE)}="款"),TRIM({x(OT_STY)}&""),"")'
    else:  # 手工
        ok = f'{x(MJ_OK)}=1'
        f[JE_DATE] = f'=IF({ok},INT({x(MJ_DATE)}),"")'
        f[JE_M] = f'=IF({ok},{x(MJ_CM)},"")'
        f[JE_MEMO] = f'=IF({ok},IF(TRIM({x(MJ_MEMO)}&"")="","手工分录",TRIM({x(MJ_MEMO)}&"")),"")'
        f[JE_DR] = f'=IF({ok},{x(MJ_DR)}&"","")'
        f[JE_CR] = f'=IF({ok},{x(MJ_CR)}&"","")'
        f[JE_AMT] = f'=IF({ok},N({x(MJ_AMT)}),0)'
        f[JE_UNIT] = f'=IF({ok},TRIM({x(MJ_UNIT)}&""),"")'
        f[JE_FEE] = f'=IF({ok},TRIM({x(MJ_FEE)}&""),"")'
    f[JE_SROW] = f'=IF({JE_VALID}{r}=1,ROW({x("A")}),"")'
    return f


def _sys_row(name, j, r):
    """系统块第 j 个槽位（0 起）在本表第 r 行的公式"""
    f = {}
    J = f'{JE_AMT}{r}'
    has = f'{J}<>0'
    if name == '收入':
        m, i = j // (JE_NCUST + 1) + 1, j % (JE_NCUST + 1)
        if i < JE_NCUST:
            cu = f'TRIM({q(SH_AUX)}!$A${AX_CUST0 + i}&"")'
            f[JE_AMT] = (f'=IF({cu}="",0,ROUND(SUMIFS({odr(OD_AMT)},{odr(OD_CUSTK)},{esc(cu)},'
                         f'{odr(OD_DM)},{m},{odr(OD_OK)},1),2))')
            f[JE_UNIT] = f'=IF({has},{cu},"")'
        else:
            f[JE_AMT] = (f'=ROUND(SUMIFS({odr(OD_AMT)},{odr(OD_DM)},{m},{odr(OD_OK)},1)'
                         f'-SUM({JE_AMT}{r - JE_NCUST}:{JE_AMT}{r - 1}),2)')
            f[JE_UNIT] = f'=IF({has},"（未登记客户）","")'
        memo, dr, cr = f'"{m}月 交货收入"', '1122', '5001'
    elif name == '工资':
        m, k = j // 4 + 1, j % 4
        dr, dest = WAGE_K[k]
        f[JE_AMT] = f'=ROUND(SUMIFS({wgr(WG_AMT)},{wgr(WG_DRC)},"{dr}",{wgr(WG_CM)},{m},{wgr(WG_OK)},1),2)'
        f[JE_FEE] = f'=IF({has},"工资","")'
        memo, cr = f'"{m}月 计提工资（{dest}）"', '2211'
    elif name == '折旧':
        m, k = j // 3 + 1, j % 3
        use = DEP_K[k]
        dr = FA_USES[use]
        f[JE_AMT] = f'=ROUND({fa_dep(m, use)},2)'
        f[JE_FEE] = f'=IF({has},"折旧费","")'
        memo, cr = f'"{m}月 计提折旧（{use}）"', '1602'
    elif name == '制造转入':
        m = j + 1
        f[JE_AMT] = f"=ROUND(N({al('制造', '本月发生', m)}),2)"
        memo, dr, cr = f'"{m}月 制造费用转入生产成本"', COMP_CODE['制造'], MOH
    else:  # 成本结转
        m, k = j // 4 + 1, j % 4
        c = COMPS[k]
        f[JE_AMT] = f"=ROUND(N({al(c, '本月转出合计', m)}),2)"
        memo, dr, cr = f'"{m}月 结转主营业务成本（{c}）"', COGS_CODE[c], COMP_CODE[c]
    f[JE_DATE] = f'=IF({has},{month_end(YR, m)},"")'
    f[JE_M] = f'=IF({has},{m},"")'
    f[JE_MEMO] = f'=IF({has},{memo},"")'
    f[JE_DR] = f'=IF({has},"{dr}","")'
    f[JE_CR] = f'=IF({has},"{cr}","")'
    return f


def build_je(wb, ctx):
    ws = wb[SH_JE]
    widths(ws, {JE_SRC: 8, JE_SROW: 7, JE_DATE: 11, JE_M: 6, JE_MEMO: 38, JE_DR: 9, JE_DRN: 18, JE_CR: 9, JE_CRN: 18,
                JE_AMT: 13, JE_UNIT: 16, JE_FEE: 10, JE_STY: 10, 'N': 2})
    title(ws, '记 账 分 录（全自动 · 一借一贷 · 所有报表都从这里取数）', JE_STY, C_RPT,
          '💡 这张表全自动，不用填、也改不了：资金日记账、送货单、外发加工、手工分录里能记账的每一行（校验不是 ✗）各生成一条分录；'
          '每月末系统再自动做几条：交货收入（按客户）、计提工资、计提折旧、制造费用转入生产成本、结转主营业务成本（按【成本分摊表】）。'
          '空着的行是没用到的槽位，不用管。用表头的筛选按「月份」「来源」「借方科目」看；想看某个科目逐笔怎么来的，去【明细账】选科目。'
          '金额是负数的是红字（退货、冲回）。「源行」是这一条在来源表里的行号，回去找原始记录用。')
    # 第 3 行
    vr, ar = jer(JE_VALID), jer(JE_AMT)
    _kpi(ws, 'A3', '分录条数', 'C3', f'=SUM({vr})', '0" 条"')
    ws.merge_cells('A3:B3')
    _kpi(ws, 'D3', '借方金额合计', 'E3', f'=SUMIFS({ar},{vr},1)', MONEY)
    put(ws, 'F3', '每条分录借贷金额相同（一借一贷），借方合计＝贷方合计。各月数、各科目数看【科目余额表】；报表看【利润表】【资产负债表】【现金流量表】。',
        F_NOTE, KPI_FILL, align=ALW)
    ws.merge_cells(f'F3:{JE_STY}3')
    ws.row_dimensions[3].height = 30
    heads = [(JE_SRC, '来源'), (JE_SROW, '源行'), (JE_DATE, '日期'), (JE_M, '月份'), (JE_MEMO, '摘要'), (JE_DR, '借方\n科目'),
             (JE_DRN, '借方名称'), (JE_CR, '贷方\n科目'), (JE_CRN, '贷方名称'), (JE_AMT, '金额'), (JE_UNIT, '往来单位'),
             (JE_FEE, '费用项目'), (JE_STY, '款式')]
    header(ws, JE_HDR, heads, HDR)
    hidden = [(JE_VALID, 'VALID 有效'), (JE_GLK, 'GLK 明细账键'), (JE_SSK, 'SSK 供应商对账键'), (JE_SLOT, 'SLOT 槽位'),
              (JE_CSK, 'CSK 客户对账键')]
    for c, t in hidden:
        ws[f'{c}{JE_HDR}'] = t
        ws[f'{c}{JE_HDR}'].font = F_HELP

    # 选择格（别的表上的）
    gsel, gun = f'TRIM({GL_CODE}&"")', f'TRIM({GL_UNIT}&"")'
    gm1, gm2 = f'N({GL_M1})', f'IF(N({GL_M2})=0,12,N({GL_M2}))'
    ssu, csu = f'TRIM({SS_UNIT}&"")', f'TRIM({CS_UNIT}&"")'

    for name, s, n in JE_BLOCKS:
        for j in range(n):
            r = s + j
            if name in SRC:
                f = _src_row(name, SRC[name][1] + j, r)
            else:
                f = _sys_row(name, j, r)
            C, D, F_, H_, J, K, O, R_ = (f'{c}{r}' for c in (JE_DATE, JE_M, JE_DR, JE_CR, JE_AMT, JE_UNIT, JE_VALID, JE_SLOT))
            f[JE_SRC] = f'=IF({O}=1,"{name}","")'
            f[JE_DRN] = f'=IF({F_}="","",IFERROR(INDEX({COA_NAMES_R},MATCH({F_},{COA_CODES},0))&"","？"&{F_}))'
            f[JE_CRN] = f'=IF({H_}="","",IFERROR(INDEX({COA_NAMES_R},MATCH({H_},{COA_CODES},0))&"","？"&{H_}))'
            f[JE_VALID] = f'=IF(ISNUMBER({D}),IF(AND({D}>=1,{D}<=12,{J}<>0),1,0),0)'
            f[JE_SLOT] = f'=ROW()-{JE_HDR}'
            key = f'{C}*100000+{R_}'
            f[JE_GLK] = (f'=IF({O}<>1,"",IF(AND({gsel}<>"",OR(LEFT({F_},LEN({gsel}))={gsel},LEFT({H_},LEN({gsel}))={gsel}),'
                         f'OR({gun}="",{K}={gun}),{D}>={gm1},{D}<={gm2}),{key},""))')
            f[JE_SSK] = (f'=IF({O}<>1,"",IF(AND({K}<>"",{K}={ssu},OR({F_}="2202",{H_}="2202"),'
                         f'{C}>=N({SS_D1}),OR(N({SS_D2})=0,{C}<=N({SS_D2}))),{key},""))')
            if name == '收入':
                f[JE_CSK] = '=""'
            else:
                f[JE_CSK] = (f'=IF({O}<>1,"",IF(AND({JE_SRC}{r}<>"收入",{K}<>"",{K}={csu},OR({F_}="1122",{H_}="1122"),'
                             f'{C}>=N({CS_D1}),OR(N({CS_D2})=0,{C}<=N({CS_D2}))),{key},""))')
            for col, v in f.items():
                ws[f'{col}{r}'] = v

    cols = [JE_SRC, JE_SROW, JE_DATE, JE_M, JE_MEMO, JE_DR, JE_DRN, JE_CR, JE_CRN, JE_AMT, JE_UNIT, JE_FEE, JE_STY]
    style_rows(ws, JE_R0, JE_R1, cols, auto=cols,
               fmts={JE_SROW: '0', JE_DATE: DATE, JE_M: MONTHF, JE_AMT: MONEY, JE_DR: '@', JE_CR: '@'},
               aligns={JE_MEMO: AL, JE_DRN: AL, JE_CRN: AL, JE_UNIT: AL, JE_AMT: AR})
    for r in range(JE_R0, JE_R1 + 1):
        for c, _ in hidden:
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in hidden])
    ws.conditional_formatting.add(f'{JE_DRN}{JE_R0}:{JE_DRN}{JE_R1}', FormulaRule(
        formula=[f'LEFT(${JE_DRN}{JE_R0},1)="？"'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(f'{JE_CRN}{JE_R0}:{JE_CRN}{JE_R1}', FormulaRule(
        formula=[f'LEFT(${JE_CRN}{JE_R0},1)="？"'], fill=FILL_WARN, font=F_RED))
    ws.auto_filter.ref = f'{JE_SRC}{JE_HDR}:{JE_STY}{JE_R1}'
    ws.freeze_panes = f'{JE_MEMO}{JE_R0}'
    print_setup(ws, f'{JE_HDR}:{JE_HDR}')


# ═══════════════════════════ 科目余额表 ═══════════════════════════
def build_tb(wb, ctx):
    ws = wb[SH_TB]
    widths(ws, {'A': 10, 'B': 26, 'C': 6, 'D': 14, 'E': 14, 'F': 14, 'G': 14, 'H': 14, 'I': 14, 'J': 14, 'K': 2})
    title(ws, '科 目 余 额 表（选月份 · 各科目年初、月初、本月发生、月末、本年累计）', 'J', C_RPT,
          '💡 C3 选月份（黄格）。每个科目一行，跟【会计科目表】一一对应；上级科目（加粗的一级科目、1002 银行存款、4001 生产成本、5401 主营业务成本）自动包含下级。'
          '余额按科目方向算：借方科目＝年初＋借−贷，贷方科目＝年初＋贷−借；负数＝跟正常方向相反（比如应付账款是负数＝多付了，变成预付）。'
          '损益类（收入、成本、费用）不结转（表结法），本年累计就是今年到所选月的发生额，利润看【利润表】。最下面一行核对借贷平不平。')
    selector(ws, 'B3', '选月份', TB_M.split('!')[1].replace('$', ''), 9, f'={AX_M}', '0"月"', '看哪个月（1～12）')
    m = TB_M
    _kpi(ws, 'D3', '期初平不平', 'E3', f"={q(SH_COA)}!$C$3")
    _kpi(ws, 'F3', '本月借贷', 'G3', f'=$C${COA_R1 + 4}')        # 最下面「核对」行的结论
    _good_bad(ws, 'E3')
    _good_bad(ws, 'G3')
    ws.merge_cells('G3:H3')
    ws.row_dimensions[3].height = 26
    heads = [('A', '科目编码'), ('B', '科目名称'), ('C', '方向'), ('D', '年初余额'), ('E', '月初余额'), ('F', '本月借方'),
             ('G', '本月贷方'), ('H', '月末余额'), ('I', '本年借方累计\n（到本月）'), ('J', '本年贷方累计\n（到本月）')]
    header(ws, 4, heads, HDR)
    ws['L4'], ws['M4'] = '上级编码', '末级'
    ws['L4'].font = ws['M4'].font = F_HELP
    assert COA_R0 == 5
    for r in range(COA_R0, COA_R1 + 1):
        src = lambda c: f'{q(SH_COA)}!{c}{r}'
        A = f'A{r}'
        f = {}
        f['A'] = f'=IF(TRIM({src(COA_CODE)}&"")="","",TRIM({src(COA_CODE)}&""))'
        f['B'] = f'=IF({A}="","",{src(COA_NAME)}&"")'
        f['C'] = f'=IF({A}="","",IF(TRIM({src(COA_DIR)}&"")="贷","贷","借"))'
        f['D'] = f'=IF({A}="","",N({src(COA_OPEN)}))'
        pre, ytd = A + '&"*"', '"<="&' + m
        f['I'] = f'=IF({A}="","",{je_sum(pre, "D", ytd)})'
        f['J'] = f'=IF({A}="","",{je_sum(pre, "C", ytd)})'
        f['F'] = f'=IF({A}="","",{je_sum(pre, "D", m)})'
        f['G'] = f'=IF({A}="","",{je_sum(pre, "C", m)})'
        f['E'] = f'=IF({A}="","",ROUND(IF(C{r}="贷",D{r}+(J{r}-G{r})-(I{r}-F{r}),D{r}+(I{r}-F{r})-(J{r}-G{r})),2))'
        f['H'] = f'=IF({A}="","",ROUND(IF(C{r}="贷",E{r}+G{r}-F{r},E{r}+F{r}-G{r}),2))'
        f['L'] = f'=TRIM({src(COA_UP)}&"")'
        f['M'] = f'=IF({A}="",0,N({src(COA_LEAF)}))'
        for col, v in f.items():
            ws[f'{col}{r}'] = v
        ws[f'L{r}'].font = ws[f'M{r}'].font = F_HELP
    cols = list('ABCDEFGHIJ')
    style_rows(ws, COA_R0, COA_R1, cols, auto=cols, fmts={c: MONEY for c in 'DEFGHIJ'} | {'A': '@'},
               aligns={'A': AL, 'B': AL, **{c: AR for c in 'DEFGHIJ'}})
    hide(ws, 'L', 'M')
    ws.conditional_formatting.add(f'A{COA_R0}:J{COA_R1}', FormulaRule(
        formula=[f'AND($A{COA_R0}<>"",$L{COA_R0}="")'], font=Font(name=YH, sz=10, bold=True, color='FF1F3864')))
    ws.conditional_formatting.add(f'D{COA_R0}:J{COA_R1}', FormulaRule(
        formula=[f'AND(ISNUMBER(D{COA_R0}),D{COA_R0}<0)'], font=F_RED))
    # 合计 / 核对
    R1, R2, R3 = COA_R1 + 2, COA_R1 + 3, COA_R1 + 4          # 86, 87, 88
    lf = f'$M${COA_R0}:$M${COA_R1}'
    put(ws, f'A{R1}', '末级科目合计（借方 / 贷方发生）', F_TXTB, FILL_SUB, align=AL)
    ws.merge_cells(f'A{R1}:E{R1}')
    for c in 'FGIJ':
        put(ws, f'{c}{R1}', f'=SUMIFS({c}${COA_R0}:{c}${COA_R1},{lf},1)', F_AUTOB, FILL_SUB, MONEY, AR)
    put(ws, f'H{R1}', '', F_AUTOB, FILL_SUB)
    put(ws, f'A{R2}', '末级科目月末余额：借方向合计 / 贷方向合计', F_TXTB, FILL_SUB, align=AL)
    ws.merge_cells(f'A{R2}:E{R2}')
    put(ws, f'F{R2}', f'=SUMIFS($H${COA_R0}:$H${COA_R1},{lf},1,$C${COA_R0}:$C${COA_R1},"借")', F_AUTOB, FILL_SUB, MONEY, AR)
    put(ws, f'G{R2}', f'=SUMIFS($H${COA_R0}:$H${COA_R1},{lf},1,$C${COA_R0}:$C${COA_R1},"贷")', F_AUTOB, FILL_SUB, MONEY, AR)
    tot_m = f'SUMIFS({jer(JE_AMT)},{jer(JE_M)},{m},{jer(JE_VALID)},1)'
    put(ws, f'A{R3}', '核对', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'B{R3}', '本月：末级借方＝末级贷方＝记账分录本月合计；月末：借方向余额＝贷方向余额', F_NOTE, FILL_TOT, align=ALW)
    chk = (f'=IF(ROUND(F{R1}-G{R1},2)<>0,"✗ 本月借贷不平：差 "&TEXT(F{R1}-G{R1},"#,##0.00"),'
           f'IF(ROUND(F{R1}-{tot_m},2)<>0,"✗ 有分录用了【会计科目表】里没有的末级科目：差 "&TEXT({tot_m}-F{R1},"#,##0.00"),'
           f'IF(ROUND(F{R2}-G{R2},2)<>0,"✗ 月末余额不平：借方向多 "&TEXT(F{R2}-G{R2},"#,##0.00")&"（先看【会计科目表】期初平不平）","√ 平")))')
    put(ws, f'C{R3}', chk, F_AUTOB, FILL_TOT, align=AL)
    ws.merge_cells(f'C{R3}:J{R3}')
    _good_bad(ws, f'C{R3}')
    ws.row_dimensions[R3].height = 22
    ws.freeze_panes = 'C5'
    print_setup(ws, '4:4', landscape=True)


# ═══════════════════════════ 明细账 ═══════════════════════════
GL_N = 500                       # 最多列 500 笔
GL_R_OPEN, GL_R0 = 5, 6          # 期初行、逐笔第一行
GL_R1 = GL_R0 + GL_N - 1         # 505
GL_R_SUM, GL_R_END = GL_R1 + 1, GL_R1 + 2
GL_SLOTC, GL_BALC = 'N', 'O'     # 隐藏：槽位序号、带符号余额（借正贷负）


def build_gl(wb, ctx):
    ws = wb[SH_GL]
    widths(ws, {'A': 6, 'B': 11, 'C': 8, 'D': 36, 'E': 26, 'F': 16, 'G': 13, 'H': 13, 'I': 6, 'J': 13, 'K': 10, 'L': 10,
                'M': 2})
    title(ws, '明 细 账（选科目看逐笔：期初 → 每一笔 → 期末）', 'L', C_RPT,
          '💡 第 3 行黄格：选科目编码（上级科目含下级，比如 1002 含所有银行），往来单位可以空着（空＝全部单位；应付 2202、应收 1122、其他应收 1221、'
          '其他应付 2241 选了单位，期初就用这一家的期初），再选起止月。下面先是期初余额，然后按日期一笔一笔列出来（最多 500 笔，多了缩小月份范围），'
          '最后是本期合计和期末余额。「方向」是余额在借方还是贷方：资产、成本、费用一般在借方；负债、权益、收入一般在贷方。')
    selector(ws, 'B3', '科目编码', 'C3', '2202', f'={COA_CODES}', '@', '【会计科目表】的科目编码；上级科目含下级')
    put(ws, 'E3', '往来单位', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'F3', None, F_SEL, FILL_SEL, align=AC)
    dv_list(ws, 'F3', f'={UN_NAMES}', '空着＝全部单位', stop=False)
    selector(ws, 'H3', '起始月', 'I3', 1, f'={AX_M}', '0', '1～12')
    selector(ws, 'J3', '截止月', 'K3', 12, f'={AX_M}', '0', '1～12')
    assert (GL_CODE, GL_UNIT, GL_M1, GL_M2) == tuple(cell(SH_GL, x) for x in ('C3', 'F3', 'I3', 'K3'))
    sel, un = 'TRIM($C$3&"")', 'TRIM($F$3&"")'
    m1, m2 = 'N($I$3)', 'IF(N($K$3)=0,12,N($K$3))'
    cnt = f'${GL_SLOTC}$4'
    ws[cnt.replace('$', '')] = f'=COUNT({jer(JE_GLK)})'
    ws[cnt.replace('$', '')].font = F_HELP
    crow = f'MATCH({sel},{COA_CODES},0)'
    put(ws, 'D3', f'=IF({sel}="","← 先选科目编码",IFERROR(INDEX({COA_NAMES_R},{crow})&"（"&IF(TRIM(INDEX({COA_DIRS},{crow})&"")="贷","贷","借")&"方科目）",'
                  f'"？【会计科目表】里没有这个编码"))', F_AUTOB, KPI_FILL, align=AC)
    put(ws, 'G3', '空＝全部单位', F_NOTE, align=AL, border=False)
    put(ws, 'L3', f'="共 "&{cnt}&" 笔"&IF({cnt}>{GL_N},"（只列前 {GL_N} 笔）","")', F_KPI_L, KPI_FILL, align=ACW)
    ws.row_dimensions[3].height = 28
    heads = [('A', '序号'), ('B', '日期'), ('C', '来源'), ('D', '摘要'), ('E', '对方科目'), ('F', '往来单位'), ('G', '借方'),
             ('H', '贷方'), ('I', '方向'), ('J', '余额'), ('K', '费用项目'), ('L', '款式')]
    header(ws, 4, heads, HDR)

    # 期初：年初（选了单位且是往来科目 → 这家的期初）＋ 起始月以前的发生
    def jsum(side, mcrit):
        col = JE_DR if side == 'D' else JE_CR
        base = f'SUMIFS({jer(JE_AMT)},{jer(col)},{sel}&"*",{jer(JE_M)},{mcrit}'
        return f'IF({un}="",{base}),{base},{jer(JE_UNIT)},{esc(un)}))'
    dirc = f'IFERROR(TRIM(INDEX({COA_DIRS},{crow})&""),"借")'
    open_coa = f'IF({dirc}="贷",-1,1)*SUMIF({COA_CODES},{sel},{COA_OPENS})'
    open_un = (f'IF({sel}="2202",-SUMIF({UN_NAMES},{esc(un)},{unr(UN_AP0)}),IF({sel}="1122",SUMIF({UN_NAMES},{esc(un)},{unr(UN_AR0)}),'
               f'IF({sel}="1221",SUMIF({UN_NAMES},{esc(un)},{unr(UN_OR0)}),IF({sel}="2241",-SUMIF({UN_NAMES},{esc(un)},{unr(UN_OP0)}),0))))')
    bo = f'{GL_BALC}{GL_R_OPEN}'
    ws[bo] = (f'=IF({sel}="",0,ROUND(IF({un}="",{open_coa},{open_un})'
              f'+{jsum("D", chr(34) + "<" + chr(34) + "&" + m1)}-{jsum("C", chr(34) + "<" + chr(34) + "&" + m1)},2))')
    ws[bo].font = F_HELP
    r = GL_R_OPEN
    put(ws, f'D{r}', f'=IF({sel}="","",IF({m1}<=1,"年初余额（建账日）","期初余额（"&{m1}&" 月初）"))', F_TXTB, FILL_SUB, align=AL)
    put(ws, f'I{r}', f'=IF({sel}="","",IF(ROUND({bo},2)>0,"借",IF(ROUND({bo},2)<0,"贷","平")))', F_TXTB, FILL_SUB, align=AC)
    put(ws, f'J{r}', f'=IF({sel}="","",ABS({bo}))', F_AUTOB, FILL_SUB, MONEY, AR)
    for c in 'ABCEFGHKL':
        put(ws, f'{c}{r}', None, F_TXT, FILL_SUB)

    # 逐笔
    for r in range(GL_R0, GL_R1 + 1):
        S = f'{GL_SLOTC}{r}'
        ix = lambda col: f'INDEX({jer(col)},{S})'
        e = f'{S}=0'
        f = {}
        f[GL_SLOTC] = f'=IF(ROW()-{GL_R_OPEN}>{cnt},0,MOD(SMALL({jer(JE_GLK)},ROW()-{GL_R_OPEN}),100000))'
        f['A'] = f'=IF({e},"",ROW()-{GL_R_OPEN})'
        f['B'] = f'=IF({e},"",{ix(JE_DATE)})'
        f['C'] = f'=IF({e},"",{ix(JE_SRC)}&"")'
        f['D'] = f'=IF({e},"",{ix(JE_MEMO)}&"")'
        drm = f'LEFT({ix(JE_DR)}&"",LEN({sel}))={sel}'
        crm = f'LEFT({ix(JE_CR)}&"",LEN({sel}))={sel}'
        f['E'] = f'=IF({e},"",IF({drm},{ix(JE_CR)}&" "&{ix(JE_CRN)},{ix(JE_DR)}&" "&{ix(JE_DRN)}))'
        f['F'] = f'=IF({e},"",{ix(JE_UNIT)}&"")'
        f['G'] = f'=IF({e},"",IF({drm},{ix(JE_AMT)},0))'
        f['H'] = f'=IF({e},"",IF({crm},{ix(JE_AMT)},0))'
        f[GL_BALC] = f'=IF({e},"",ROUND(N({GL_BALC}{r - 1})+G{r}-H{r},2))'
        f['I'] = f'=IF({e},"",IF({GL_BALC}{r}>0,"借",IF({GL_BALC}{r}<0,"贷","平")))'
        f['J'] = f'=IF({e},"",ABS({GL_BALC}{r}))'
        f['K'] = f'=IF({e},"",{ix(JE_FEE)}&"")'
        f['L'] = f'=IF({e},"",{ix(JE_STY)}&"")'
        for col, v in f.items():
            ws[f'{col}{r}'] = v
        ws[S].font = ws[f'{GL_BALC}{r}'].font = F_HELP
    cols = list('ABCDEFGHIJKL')
    style_rows(ws, GL_R0, GL_R1, cols, auto=cols, fmts={'B': DATE, 'G': MONEY, 'H': MONEY, 'J': MONEY, 'A': INT},
               aligns={'D': AL, 'E': AL, 'F': AL, 'G': AR, 'H': AR, 'J': AR})
    # 合计、期末
    r = GL_R_SUM
    put(ws, f'D{r}', f'="本期合计（"&{m1}&"～"&{m2}&" 月）"', F_TXTB, FILL_TOT, align=AL)
    mc1, mc2 = f'">="&{m1}', f'"<="&{m2}'

    def jsum2(side):
        col = JE_DR if side == 'D' else JE_CR
        base = f'SUMIFS({jer(JE_AMT)},{jer(col)},{sel}&"*",{jer(JE_M)},{mc1},{jer(JE_M)},{mc2}'
        return f'IF({sel}="",0,IF({un}="",{base}),{base},{jer(JE_UNIT)},{esc(un)})))'
    put(ws, f'G{r}', f'={jsum2("D")}', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, f'H{r}', f'={jsum2("C")}', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'ABCEFIJKL':
        put(ws, f'{c}{r}', None, F_TXT, FILL_TOT)
    re_ = GL_R_END
    be = f'{GL_BALC}{re_}'
    ws[be] = f'=ROUND({bo}+G{r}-H{r},2)'
    ws[be].font = F_HELP
    put(ws, f'D{re_}', '期末余额', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'I{re_}', f'=IF({sel}="","",IF({be}>0,"借",IF({be}<0,"贷","平")))', F_TXTB, FILL_TOT, align=AC)
    put(ws, f'J{re_}', f'=IF({sel}="","",ABS({be}))', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'ABCEFGHKL':
        put(ws, f'{c}{re_}', None, F_TXT, FILL_TOT)
    put(ws, f'A{re_ + 1}', '期末余额＝期初＋本期借方−本期贷方（借正贷负）；「方向」借＝余额在借方，贷＝在贷方。超过 500 笔时下面只列前 500 笔，但合计、期末是全部的。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{re_ + 1}:L{re_ + 1}')
    hide(ws, GL_SLOTC, GL_BALC)
    ws.conditional_formatting.add(f'G{GL_R0}:H{GL_R1}', FormulaRule(formula=[f'AND(ISNUMBER(G{GL_R0}),G{GL_R0}=0)'],
                                                                     font=Font(name=YH, sz=10, color='FFD9D9D9')))
    ws.freeze_panes = 'A5'
    print_setup(ws, '4:4', landscape=True)


# ═══════════════════════════ 利润表 ═══════════════════════════
MC = [CL(2 + i) for i in range(12)]          # B..M ＝ 1～12 月
C_YTD, C_PCT = 'N', 'O'


def _pl(code, m, credit):
    """某科目（前缀，含下级）第 m 月净发生：credit=True 贷−借，否则 借−贷"""
    d, c = je_sum(f'"{code}*"', 'D', str(m)), je_sum(f'"{code}*"', 'C', str(m))
    return f'({c}-{d})' if credit else f'({d}-{c})'


def build_is(wb, ctx):
    ws = wb[SH_IS]
    widths(ws, {'A': 30, **{c: 12 for c in MC}, C_YTD: 14, C_PCT: 9})
    title(ws, '利 润 表（1～12 月 ＋ 全年 · 表结法）', C_PCT, C_RPT,
          '💡 全自动。收入按交货月份（【订单明细】交货日期）；营业成本＝每月按【成本分摊表】结转的「本月转出合计」（只算交了货的鞋的成本，没交货的挂在生产成本里）；'
          '费用按付款 / 计提的月份。晚到的送货单录进去以后，以前月份的成本、利润会跟着变（除非在【基础资料】设了已结账月份）。'
          '下面几行：毛利率、净利率、交货双数、平均每双成本 / 售价。内账含税口径，不拆增值税。')
    put(ws, 'A3', f'={CO}&"　"&{YR}&" 年度（单位：元）"', F_KPI_L, align=AL, border=False)
    header(ws, 4, [('A', '项  目')] + [(c, f'{i + 1}月') for i, c in enumerate(MC)] + [(C_YTD, '全年'), (C_PCT, '占收入\n（全年）')],
           HDR)
    rows = [  # (键, 标签, 类型, 公式/参数, 级别)
        ('rev', '一、营业收入', 'pl', [('5001', True), ('5051', True)], 0),
        ('cogs', '减：营业成本', 'pl', [('5401', False), ('5402', False)], 1),
        ('c1', '　　其中：材料', 'pl', [(COGS_CODE['材料'], False)], 2),
        ('c2', '　　　　　外发加工', 'pl', [(COGS_CODE['外发'], False)], 2),
        ('c3', '　　　　　人工', 'pl', [(COGS_CODE['人工'], False)], 2),
        ('c4', '　　　　　制造费用', 'pl', [(COGS_CODE['制造'], False)], 2),
        ('tax', '　　税金及附加', 'pl', [('5403', False)], 1),
        ('sell', '　　销售费用', 'pl', [('5601', False)], 1),
        ('adm', '　　管理费用', 'pl', [('5602', False)], 1),
        ('fin', '　　财务费用', 'pl', [('5603', False)], 1),
        ('op', '二、营业利润', 'calc', 'rev-cogs-tax-sell-adm-fin', 0),
        ('nii', '加：营业外收入', 'pl', [('5301', True)], 1),
        ('nie', '减：营业外支出', 'pl', [('5711', False)], 1),
        ('ebt', '三、利润总额', 'calc', 'op+nii-nie', 0),
        ('itx', '减：所得税费用', 'pl', [('5801', False)], 1),
        ('np', '四、净利润', 'calc', 'ebt-itx', 0),
        (None, '', 'blank', None, 0),
        ('gp', '毛利（营业收入−营业成本）', 'calc', 'rev-cogs', 1),
        ('gpr', '毛利率', 'pct', ('gp', 'rev'), 1),
        ('npr', '净利率', 'pct', ('np', 'rev'), 1),
        ('qty', '交货双数', 'qty', None, 1),
        ('ucost', '平均每双成本', 'per', ('cogs', 'qty'), 1),
        ('uprice', '平均每双售价', 'per', ('rev', 'qty'), 1),
    ]
    R = {}
    r = 5
    for key, *_x in rows:
        if key:
            R[key] = r
        r += 1
    r = 5
    for key, lab, kind, arg, lvl in rows:
        if kind == 'blank':
            r += 1
            continue
        bold = lvl == 0
        fl = FILL_TOT if bold else (FILL_SUB if kind in ('pct', 'per', 'qty') or key == 'gp' else None)
        put(ws, f'A{r}', lab, F_TXTB if lvl < 2 else F_TXT, fl, align=AL)
        for i, c in enumerate(MC + [C_YTD]):
            m = i + 1
            if kind == 'pl':
                f = '+'.join(_pl(code, m, cr) for code, cr in arg) if c != C_YTD else f'SUM({MC[0]}{r}:{MC[-1]}{r})'
                f = f'ROUND({f},2)'
            elif kind == 'calc':
                f = re.sub(r'[a-z][a-z0-9]*', lambda mm: f'{c}{R[mm.group()]}', arg)
                f = f'ROUND({f},2)'
            elif kind == 'pct':
                a, b = arg
                f = f'IF(N({c}{R[b]})=0,"",ROUND({c}{R[a]}/{c}{R[b]},4))'
            elif kind == 'qty':
                f = f'SUM({smr("双数", m)})' if c != C_YTD else f'SUM({MC[0]}{r}:{MC[-1]}{r})'
            else:  # per
                a, b = arg
                f = f'IF(N({c}{R[b]})=0,"",ROUND({c}{R[a]}/{c}{R[b]},2))'
            fmt = PCT if kind == 'pct' else (INT if kind == 'qty' else MONEY)
            put(ws, f'{c}{r}', f'={f}', F_AUTOB if bold else F_AUTO, fl, fmt, AR)
        if kind in ('pl', 'calc'):
            put(ws, f'{C_PCT}{r}', f'=IF(N(${C_YTD}${R["rev"]})=0,"",ROUND({C_YTD}{r}/${C_YTD}${R["rev"]},4))', F_NOTE, fl, PCT, AC)
        else:
            put(ws, f'{C_PCT}{r}', None, F_NOTE, fl)
        r += 1
    last = r - 1
    put(ws, f'A{last + 2}', '说明：营业收入＝5001 主营业务收入＋5051 其他业务收入（贷−借）；营业成本＝5401 主营业务成本（含 540101～540104）＋5402 其他业务成本；'
                            '各费用＝对应科目的借−贷（含自己加的下级科目）。交货双数来自【_款式月】（只算实交>0 的行）。'
                            '【会计科目表】里自己加的其他损益科目（编码不在上面这些下面的）不在这张表里，但会算进【资产负债表】的本年利润。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{last + 2}:{C_PCT}{last + 3}')
    ws.row_dimensions[last + 2].height = 30
    ws.conditional_formatting.add(f'B5:{C_YTD}{last}', FormulaRule(formula=['AND(ISNUMBER(B5),B5<0)'], font=F_RED))
    ws.freeze_panes = 'B5'
    print_setup(ws, '4:4', landscape=True)


# ═══════════════════════════ 资产负债表 ═══════════════════════════
# 右边隐藏区（行＝会计科目表 80 行）：K 编码 L 类别 M 方向 N 报表项目 O 归到哪 P 年初（借正贷负） Q 所选月末（借正贷负）；U 报表行清单
BH = dict(code='K', cls='L', dir='M', line='N', key='O', open='P', end='Q')
BS_LINES_A = ['货币资金', '应收账款', '预付账款', '其他应收款', '存货', '固定资产原价', '减：累计折旧', '长期待摊费用']
BS_LINES_L = ['短期借款', '应付账款', '预收账款', '应付职工薪酬', '应交税费', '其他应付款', '实收资本', '未分配利润']
SPLIT = {'1122': (UN_BAR, UN_AR0, 1), '2202': (UN_BAP, UN_AP0, -1), '1221': (UN_BOR, UN_OR0, 1), '2241': (UN_BOP, UN_OP0, -1)}
SUM_PARENTS = {c for c, _n, _cl, _d, _u, _l, src in COA if src == 'sum'}


def build_bs(wb, ctx):
    ws = wb[SH_BS]
    widths(ws, {'A': 22, 'B': 5, 'C': 15, 'D': 15, 'E': 24, 'F': 5, 'G': 15, 'H': 15, 'I': 16, 'J': 2})
    title(ws, '资 产 负 债 表（选月份 · 期末数 ＋ 年初数）', 'H', C_RPT,
          '💡 C3 选月份（黄格），看那个月月底的家底；年初数＝建账那天的数（【会计科目表】年初余额）。H3 是「资产总计−负债和权益总计」，0 才对。'
          '应收 / 应付按每一家拆方向：供应商多付了的算「预付账款」，客户多付了的算「预收账款」。存货＝车间里还没交货的鞋的成本（生产成本余额，'
          '直接记到款式、款式还没交货的挂在这里）。未分配利润＝年初未分配利润＋今年到所选月的利润（表结法，不做结转）。')
    put(ws, 'A3', f'="编制单位："&{CO}', F_KPI_L, align=AL, border=False)
    assert BS_M == cell(SH_BS, 'C3') and BS_CHECK == cell(SH_BS, 'H3')
    selector(ws, 'B3', '月份', 'C3', 9, f'={AX_M}', '0"月"', '看哪个月月底（1～12）')
    put(ws, 'D3', f'={YR}&"年"&$C$3&"月"&DAY({month_end(YR, "$C$3")})&"日"', F_KPI_L, align=AC, border=False)
    put(ws, 'E3', '单位：元', F_NOTE, align=AC, border=False)
    put(ws, 'G3', '平不平', F_KPI_L, KPI_FILL, align=AC)
    ws.row_dimensions[3].height = 26
    header(ws, 4, [('A', '资    产'), ('B', '行次'), ('C', '期末数'), ('D', '年初数\n（建账日）'),
                   ('E', '负债和所有者权益'), ('F', '行次'), ('G', '期末数'), ('H', '年初数\n（建账日）')], HDR)

    # ── 隐藏区：每个科目的余额
    K, L, M, N, O, Pc, Q = (BH[k] for k in ('code', 'cls', 'dir', 'line', 'key', 'open', 'end'))
    for c, t in ((K, '编码'), (L, '类别'), (M, '方向'), (N, '报表项目'), (O, '归到哪'), (Pc, '年初(借正)'), (Q, '月末(借正)'), ('U', '报表行')):
        ws[f'{c}4'] = t
        ws[f'{c}4'].font = F_HELP
    lines_rng = f'$U$5:$U${4 + len(BS_LINES_A + BS_LINES_L)}'
    for i, nm in enumerate(BS_LINES_A + BS_LINES_L):
        ws[f'U{5 + i}'] = nm
        ws[f'U{5 + i}'].font = F_HELP
    preset = {COA_R0 + i: c[0] for i, c in enumerate(COA)}
    for r in range(COA_R0, COA_R1 + 1):
        src = lambda c: f'{q(SH_COA)}!{c}{r}'
        k = f'{K}{r}'
        f = {}
        f[K] = f'=IF(TRIM({src(COA_CODE)}&"")="","",TRIM({src(COA_CODE)}&""))'
        f[L] = f'=TRIM({src(COA_CLS)}&"")'
        f[M] = f'=IF(TRIM({src(COA_DIR)}&"")="贷","贷","借")'
        f[N] = f'=TRIM({src(COA_LINE)}&"")'
        sp = 'OR(' + ','.join(f'{k}="{c}"' for c in SPLIT) + ')'
        f[O] = (f'=IF({k}="","",IF({L}{r}="损益","#损益",IF({sp},"#往来",IF(AND({N}{r}<>"",COUNTIF({lines_rng},{N}{r})>0),{N}{r},'
                f'IF(OR({L}{r}="负债",{L}{r}="权益"),"#其他负债","#其他资产")))))')
        # 年初：上级（下级合计）科目自己不算，免得重复
        if preset.get(r) in SUM_PARENTS:
            f[Pc] = f'=0'
        else:
            f[Pc] = f'=IF({k}="",0,IF({M}{r}="贷",-1,1)*N({src(COA_OPEN)}))'
        mc = f'"<="&$C$3'
        f[Q] = (f'=IF({k}="",0,ROUND({Pc}{r}+{je_sum(k, "D", mc)}-{je_sum(k, "C", mc)},2))')
        for col, v in f.items():
            ws[f'{col}{r}'] = v
            ws[f'{col}{r}'].font = F_HELP
    # 记账分录里用了【会计科目表】没有的科目（不该有）：放「其他」，保证平衡
    unm_e, unm_o = f'$Q${COA_R1 + 1}', f'$P${COA_R1 + 1}'
    ws[f'{Q}{COA_R1 + 1}'] = f'=ROUND(-(SUM($Q${COA_R0}:$Q${COA_R1})-SUM($P${COA_R0}:$P${COA_R1})),2)'
    ws[f'{Pc}{COA_R1 + 1}'] = '=0'
    ws[f'{O}{COA_R1 + 1}'] = '科目表没有的'
    for c in (Q, Pc, O):
        ws[f'{c}{COA_R1 + 1}'].font = F_HELP
    hide(ws, K, L, M, N, O, Pc, Q, 'R', 'S', 'T', 'U')

    def g(line, col):
        return f'SUMIF($O${COA_R0}:$O${COA_R1},"{line}",${col}${COA_R0}:${col}${COA_R1})'

    def parts(col, which):
        """col：P 年初 / Q 期末；which：0 年初（用往来单位期初列）/ 1 期末（用 UN_Bxx 列）"""
        def T(code):            # 该科目正常方向余额
            sign = SPLIT[code][2]
            return f'{"-" if sign < 0 else ""}SUMIF($K${COA_R0}:$K${COA_R1},"{code}",${col}${COA_R0}:${col}${COA_R1})'

        def U(code):
            return unr(SPLIT[code][1 - which])

        def pos(code):
            return f'SUMIF({U(code)},">0")'

        def neg(code):
            return f'-SUMIF({U(code)},"<0")'

        def rest(code):
            return f'({T(code)}-SUM({U(code)}))'
        unm = unm_o if which == 0 else unm_e
        d = {
            '货币资金': g('货币资金', col),
            '应收账款': f'{g("应收账款", col)}+{pos("1122")}+MAX({rest("1122")},0)',
            '预付账款': f'{g("预付账款", col)}+{neg("2202")}+MAX(-{rest("2202")},0)',
            '其他应收款': f'{g("其他应收款", col)}+{pos("1221")}+MAX({rest("1221")},0)+{neg("2241")}+MAX(-{rest("2241")},0)',
            '存货': g('存货', col),
            '其他资产': f'{g("#其他资产", col)}+MAX({unm},0)',
            '固定资产原价': g('固定资产原价', col),
            '减：累计折旧': f'-{g("减：累计折旧", col)}',
            '长期待摊费用': g('长期待摊费用', col),
            '短期借款': f'-{g("短期借款", col)}',
            '应付账款': f'-{g("应付账款", col)}+{pos("2202")}+MAX({rest("2202")},0)',
            '预收账款': f'-{g("预收账款", col)}+{neg("1122")}+MAX(-{rest("1122")},0)',
            '应付职工薪酬': f'-{g("应付职工薪酬", col)}',
            '应交税费': f'-{g("应交税费", col)}',
            '其他应付款': f'-{g("其他应付款", col)}+{pos("2241")}+MAX({rest("2241")},0)+{neg("1221")}+MAX(-{rest("1221")},0)',
            '其他负债': f'-{g("#其他负债", col)}+MAX(-{unm},0)',
            '实收资本': f'-{g("实收资本", col)}',
            '本年利润': f'-{g("#损益", col)}',
        }
        d['未分配利润'] = f'-{g("未分配利润", col)}+({d["本年利润"]})'
        return d

    left = [('流动资产：', 'head'), ('货币资金', 'v'), ('应收账款', 'v'), ('预付账款', 'v'), ('其他应收款', 'v'), ('存货', 'v'),
            ('其他（未归类科目）', 'v:其他资产'), ('流动资产合计', 'sum:2-7'), ('非流动资产：', 'head'), ('固定资产原价', 'v'),
            ('减：累计折旧', 'v'), ('固定资产账面价值', 'fa'), ('长期待摊费用', 'v'), ('非流动资产合计', 'nca')]
    right = [('流动负债：', 'head'), ('短期借款', 'v'), ('应付账款', 'v'), ('预收账款', 'v'), ('应付职工薪酬', 'v'), ('应交税费', 'v'),
             ('其他应付款', 'v'), ('其他（未归类科目）', 'v:其他负债'), ('负债合计', 'sum:2-8'), ('所有者权益：', 'head'),
             ('实收资本', 'v'), ('未分配利润', 'v'), ('　其中：本年利润', 'info:本年利润'), ('所有者权益合计', 'eq')]
    assert len(left) == len(right)
    R0 = 5
    P0, P1 = parts('P', 0), parts('Q', 1)
    seq = {'A': 0, 'E': 0}

    def side(items, cl, cn, ce, co):
        rows = {}
        for i, (lab, kind) in enumerate(items):
            r = R0 + i
            rows[lab] = r
            head = kind == 'head'
            tot = kind.startswith('sum') or kind in ('nca', 'eq')
            fl = FILL_HEAD if head else (FILL_SUB if tot or kind == 'fa' else None)
            put(ws, f'{cl}{r}', lab, F_HEADB if head else (F_TXTB if tot else F_TXT), fl, align=AL)
            if head:
                for c in (cn, ce, co):
                    put(ws, f'{c}{r}', None, F_TXT, fl)
                continue
            if not kind.startswith('info'):
                seq[cl] += 1
            put(ws, f'{cn}{r}', None if kind.startswith('info') else seq[cl], F_NOTE, fl, align=AC)
            for c, d in ((ce, P1), (co, P0)):
                if kind == 'v':
                    f = d[lab]
                elif kind.startswith('v:') or kind.startswith('info:'):
                    f = d[kind.split(':')[1]]
                elif kind.startswith('sum'):
                    a, b = (int(x) for x in kind[4:].split('-'))
                    f = f'SUM({c}{R0 + a - 1}:{c}{R0 + b - 1})'
                elif kind == 'fa':
                    f = f'{c}{rows["固定资产原价"]}-{c}{rows["减：累计折旧"]}'
                elif kind == 'nca':
                    f = f'{c}{rows["固定资产账面价值"]}+{c}{rows["长期待摊费用"]}'
                else:  # eq
                    f = f'{c}{rows["实收资本"]}+{c}{rows["未分配利润"]}'
                put(ws, f'{c}{r}', f'=ROUND({f},2)', F_AUTOB if tot else (F_NOTE if kind.startswith('info') else F_AUTO), fl,
                    MONEY, AR)
        return rows
    rl = side(left, 'A', 'B', 'C', 'D')
    rr = side(right, 'E', 'F', 'G', 'H')
    rt = R0 + len(left) + 1
    for c in 'ABCDEFGH':
        put(ws, f'{c}{rt - 1}', None, F_TXT)
    put(ws, f'A{rt}', '资产总计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'B{rt}', None, F_TXT, FILL_TOT)
    put(ws, f'E{rt}', '负债和所有者权益总计', F_TXTB, FILL_TOT, align=AL)
    put(ws, f'F{rt}', None, F_TXT, FILL_TOT)
    for c in 'CD':
        put(ws, f'{c}{rt}', f'=ROUND({c}{rl["流动资产合计"]}+{c}{rl["非流动资产合计"]},2)', F_AUTOB, FILL_TOT, MONEY, AR)
    for c in 'GH':
        put(ws, f'{c}{rt}', f'=ROUND({c}{rr["负债合计"]}+{c}{rr["所有者权益合计"]},2)', F_AUTOB, FILL_TOT, MONEY, AR)
    # H3 平不平
    put(ws, 'H3', f'=ROUND(C{rt}-G{rt},2)', F_VAL, KPI_FILL, '#,##0.00;[Red]-#,##0.00;0', AC)
    put(ws, 'I3', f'=IF(H3<>0,"✗ 期末差 "&TEXT(H3,"#,##0.00"),IF(ROUND(D{rt}-H{rt},2)<>0,"✗ 年初差 "&TEXT(D{rt}-H{rt},"#,##0.00")&"（看【会计科目表】期初）","√ 平"))',
        F_VAL, KPI_FILL, align=AC)
    _good_bad(ws, 'I3')
    ws.conditional_formatting.add('H3', FormulaRule(formula=['H3<>0'], fill=FILL_WARN, font=F_BAD))
    notes = [
        '说明：',
        '· 货币资金＝库存现金＋银行存款＋其他货币资金（支付宝 / 微信），所选月份是 10 月以后时应该等于【资金日记账】的总余额（✗ 的行不进账，余额里有、这里没有）。',
        '· 应收账款 / 预收账款、应付账款 / 预付账款、其他应收款 / 其他应付款：按【往来单位】每一家的余额拆方向，没写单位的按科目余额方向。',
        '· 存货＝生产成本（4001，含直接材料、外发、人工、制造费用）＋制造费用（4101，月末转完应为 0）：都是还没交货的在制品。',
        '· 未分配利润＝利润分配（3104）年初余额＋本年利润（所有损益类科目，年初到所选月）。',
        '· 其他（未归类科目）：【会计科目表】里自己加的科目，「报表项目」跟这张表的行名对不上的放这里，平衡不受影响。',
    ]
    for i, t in enumerate(notes):
        rr_ = rt + 2 + i
        put(ws, f'A{rr_}', t, F_NOTE if i else F_TXTB, border=False, align=AL)
        ws.merge_cells(f'A{rr_}:H{rr_}')
    ws.freeze_panes = 'A5'
    print_setup(ws, '4:4', landscape=False)


# ═══════════════════════════ 现金流量表 ═══════════════════════════
def build_cf(wb, ctx):
    ws = wb[SH_CF]
    widths(ws, {'A': 34, **{c: 12 for c in MC}, C_YTD: 14})
    title(ws, '现 金 流 量 表（直接法简表 · 1～12 月 ＋ 全年）', C_YTD, C_RPT,
          '💡 全自动，来自【资金日记账】能记账的行（每一类收支在【基础资料】③ 设了算哪一项现金流量；费用支出按费用项目归类）。'
          '收进来是正数、付出去是负数；自己账户之间转钱不算。期末现金余额应该等于资金科目（库存现金、银行存款、支付宝微信）月末余额，'
          '最下面核对行不是 0，多半是【手工分录】动了资金科目（手工分录不进现金流量表）。')
    put(ws, 'A3', f'={CO}&"　"&{YR}&" 年度（单位：元）"', F_KPI_L, align=AL, border=False)
    header(ws, 4, [('A', '项  目')] + [(c, f'{i + 1}月') for i, c in enumerate(MC)] + [(C_YTD, '全年')], HDR)
    groups = [('经营', '一、经营活动产生的现金流量', '经营活动产生的现金流量净额'),
              ('投资', '二、投资活动产生的现金流量', '投资活动产生的现金流量净额'),
              ('筹资', '三、筹资活动产生的现金流量', '筹资活动产生的现金流量净额')]
    net, cm, cf, ok, cat = jr(J_NET), jr(J_CM), jr(J_CF), jr(J_OK), jr(J_CATX)
    r = 5
    nets = []
    item_rows = []
    allcols = MC + [C_YTD]

    def row(lab, fn, font=F_AUTO, fl=None, lfont=F_TXT, fmt=MONEY):
        nonlocal r
        put(ws, f'A{r}', lab, lfont, fl, align=AL)
        for i, c in enumerate(allcols):
            put(ws, f'{c}{r}', f'={fn(i + 1, c)}', font, fl, fmt, AR)
        r += 1
        return r - 1

    def ysum(c, rr):
        return f'SUM({MC[0]}{rr}:{MC[-1]}{rr})'
    # 先算好每一组「项目」行在哪几行（「其他」行要减掉全部三组的项目）
    item_ranges, pos = [], r
    for grp, _h, _t in groups:
        n_it = sum(1 for _nm, g_ in CF_ITEMS if g_ == grp)
        item_ranges.append((pos + 1, pos + n_it))
        pos += 1 + n_it + (1 if grp == '经营' else 0) + 1
    for grp, head, tot in groups:
        put(ws, f'A{r}', head, F_HEADB, FILL_HEAD, align=AL)
        for c in allcols:
            put(ws, f'{c}{r}', None, F_TXT, FILL_HEAD)
        r += 1
        first = r
        for nm, g_ in CF_ITEMS:
            if g_ != grp:
                continue
            rr = row(f'　{nm}', lambda m, c, nm=nm: ysum(c, r) if c == C_YTD else
                     f'ROUND(SUMIFS({net},{cf},"{nm}",{cm},{m},{ok},1),2)')
            item_rows.append(rr)
        if grp == '经营':
            # 收支类别没设现金流量项目（或填了清单以外的）的，放这里，保证净增加额对得上
            rr = r
            row('　其他（收支类别没设现金流量项目的）', lambda m, c: ysum(c, rr) if c == C_YTD else
                f'ROUND(SUMIFS({net},{cm},{m},{ok},1,{cat},"<>内部转账")-'
                + '-'.join(f'SUM({c}{a}:{c}{b})' for a, b in item_ranges) + ',2)', font=F_NOTE, lfont=F_NOTE)
        assert (first, first + sum(1 for _nm, g_ in CF_ITEMS if g_ == grp) - 1) in item_ranges
        last = r - 1
        nets.append(row(tot, lambda m, c: f'ROUND(SUM({c}{first}:{c}{last}),2)', F_AUTOB, FILL_SUB, F_TXTB))
    r_inc = row('四、现金净增加额', lambda m, c: '+'.join(f'{c}{x}' for x in nets), F_AUTOB, FILL_TOT, F_TXTB)
    r_beg = r
    open_all = f'SUM({AC_OPENS})'
    row('加：期初现金余额', lambda m, c: open_all if m == 1 else (f'{MC[0]}{r_beg}' if c == C_YTD else f'{MC[m - 2]}{r_beg + 1}'))
    r_end = row('五、期末现金余额', lambda m, c: f'ROUND({c}{r_inc}+{c}{r_beg},2)', F_AUTOB, FILL_TOT, F_TXTB)
    r += 1
    put(ws, f'A{r}', '核对', F_HEADB, FILL_HEAD, align=AL)
    for c in allcols:
        put(ws, f'{c}{r}', None, F_TXT, FILL_HEAD)
    r += 1
    fund_open = '+'.join(f'SUMIF({COA_CODES},"{c}",{COA_OPENS})' for c in ('1001', '1002', '1012'))

    def fund_bal(m):
        mc = f'"<="&{m}'
        return '+'.join(f'{je_sum(chr(34) + p + "*" + chr(34), "D", mc)}-{je_sum(chr(34) + p + "*" + chr(34), "C", mc)}'
                        for p in ('1001', '1002', '1012'))
    r_fund = row('资金科目月末余额（1001 / 1002 / 1012）', lambda m, c: f'ROUND({fund_open}+{fund_bal(12 if c == C_YTD else m)},2)')
    r_diff = row('差额（期末现金−资金科目，应为 0）', lambda m, c: f'ROUND({c}{r_end}-{c}{r_fund},2)', F_AUTOB)
    row('对得上吗', lambda m, c: f'IF({c}{r_diff}=0,"√","✗ 看手工分录")', F_AUTOB, fmt=None)
    ws.conditional_formatting.add(f'B{r - 1}:{C_YTD}{r - 1}', FormulaRule(formula=[f'LEFT(B{r - 1},1)="✗"'], fill=FILL_WARN, font=F_BAD))
    ws.conditional_formatting.add(f'B{r - 1}:{C_YTD}{r - 1}', FormulaRule(formula=[f'LEFT(B{r - 1},1)="√"'], font=F_GOOD))
    ws.conditional_formatting.add(f'B5:{C_YTD}{r - 2}', FormulaRule(formula=['AND(ISNUMBER(B5),B5<0)'], font=F_RED))
    put(ws, f'A{r + 1}', '说明：每一项＝【资金日记账】里能记账（校验不是 ✗）、现金流量项目是这一项、成本月份是这个月的「收入−支出」。'
                         '期初现金余额：1 月＝各资金账户期初合计（建账日），以后＝上月期末。全年列的期初＝1 月期初、期末＝12 月期末。',
        F_NOTE, border=False, align=ALW)
    ws.merge_cells(f'A{r + 1}:{C_YTD}{r + 2}')
    ws.freeze_panes = 'B5'
    print_setup(ws, '4:4', landscape=True)


def build(wb, ctx):
    build_je(wb, ctx)
    build_tb(wb, ctx)
    build_gl(wb, ctx)
    build_is(wb, ctx)
    build_bs(wb, ctx)
    build_cf(wb, ctx)
