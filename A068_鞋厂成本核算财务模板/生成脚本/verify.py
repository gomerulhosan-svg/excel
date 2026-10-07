# -*- coding: utf-8 -*-
"""独立复算：不看表里的公式，直接从各录入表（资金日记账、送货单、外发、订单、工资、手工分录）和档案的「录入格」用 Python 按同样的口径重算，
   再跟 LibreOffice 算出来的表里的数逐项比（款式×月 的双数/收入/未交/转出/成本、成本分摊表、订单成本、订单汇总、利润表、科目余额表、
   现金流量表、资产负债表、账户余额、往来单位余额、首页往来、款式档案「挂着没转」……）。

口径（第二轮）：
  · 送货单 / 外发的成本月：【基础资料】⑩ 每月结账日期——送货月填了结账日期、收单日期（空＝送货日期）晚于它 → 收单月；否则送货月。
  · 订单「成本双数」CQ：实交>0 且备注不含「返修」；退货（负数）、返修重交只算收入，不算双数、不分成本。
  · 直接成本按交货进度转出：转X＝ROUND((到本月累计直接成本−以前已转出)×双数÷(双数＋月末未交订单双数),2)；
    未交＝下单日期<下月1日（空也算）的订单数量 − 交货日期<下月1日的成本双数。
  · 停产：款式档案状态＝停产 → REL＝最后一个有双数的月份＋1（一双没交过＝1）；m≥REL 时这个款还挂着的直接成本当月转进公共（_款式月 停X、成本分摊表 停产转入）。
  · 资金日记账金额：粘贴来的文本数字（「12,000.00」「¥50」）照 common.num() 转。
  · 「折旧提到几月」空着＝最新月份 LASTM＝各录入表能记账的行里最大的月份（日记账、送货单 CM、外发 CM、订单 DM、工资、手工）。
  · 年初在制（会计科目表 400101～400104「年初余额（用的）」H 列）进 1 月「公共上月结转」。

跑法：python3 verify.py <LibreOffice 算过的副本.xlsx>      （build.py --keep 留下的 _calc/…_算数.xlsx）
"""
import datetime as dt
import re
import sys
from collections import defaultdict

import openpyxl
from layout import *

TOL = 0.011
EPOCH = dt.date(1899, 12, 30)
XFER = '内部转账'
SUPS = ('材料供应商', '外发加工厂')
COMP_OF = {COMP_CODE['材料']: '材料', COMP_CODE['外发']: '外发', COMP_CODE['人工']: '人工', MOH: '制造'}
SYS_CODES = {COMP_CODE['制造']} | {COGS_CODE[c] for c in COMPS}
CF_BORROW, CF_REPAY = '借款收到的现金', '偿还借款支付的现金'


# ─────────────────────────── 小工具：照 Excel 的规矩 ───────────────────────────
def d2(x):
    return round(x + 1e-9 if x >= 0 else x - 1e-9, 2)


def d4(x):
    return round(x + 1e-11 if x >= 0 else x - 1e-11, 4)


def isnum(v):
    """ISNUMBER：数字或日期"""
    return isinstance(v, (int, float, dt.date, dt.datetime)) and not isinstance(v, bool)


def as_date(v):
    """日期格 / 日期序号 → date；别的 None"""
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        try:
            return EPOCH + dt.timedelta(days=int(v))
        except (OverflowError, ValueError):
            return None
    return None


def serial(v):
    """数字或日期 → Excel 序号（日期带时刻的含小数）"""
    if isinstance(v, dt.datetime):
        return (v.date() - EPOCH).days + (v.hour * 3600 + v.minute * 60 + v.second) / 86400
    if isinstance(v, dt.date):
        return (v - EPOCH).days
    return float(v)


def N(v):
    """Excel N()：数字原样、日期取序号、别的 0"""
    return serial(v) if isnum(v) else 0.0


def blank(v):
    """TRIM(x&"")="" """
    return v is None or (isinstance(v, str) and xtrim(v) == '')


def numv(v):
    """数量 / 单价 / 金额格：空 0、数字原样、文字 None（这一行 ✗）"""
    if blank(v):
        return 0.0
    if isnum(v):
        return N(v)
    return None


def txt2num(v):
    """资金日记账金额：真数字原样；文本数字（「12,000.00」「¥50」「 50 」）照 common.num() 转；转不出来 None"""
    if v is None or v == '':
        return 0.0
    if isnum(v):
        return N(v)
    s = xtrim(str(v))
    if s == '':
        return 0.0
    s = s.replace('¥', '').replace('￥', '').replace(',', '')
    try:
        return float(s)
    except ValueError:
        return None


def xtrim(s):
    """Excel TRIM：去掉头尾空格、中间连着的空格并成一个（只管半角空格）"""
    return re.sub(' +', ' ', s).strip(' ')


def xtext(v):
    """Excel 的 x&"" """
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, (dt.date, dt.datetime)):
        v = serial(v)
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else f'{v:.15g}'
    return str(v)


def tkey(v):
    """TRIM(x&"") 再按 Excel 比较不分大小写"""
    return xtrim(xtext(v)).casefold()


def cell_of(addr):
    return addr.split('!')[1].replace('$', '')


def grab(ws, r0, r1, cols):
    """一次读一块：[(行号, {列: 值})]"""
    idx = {c: openpyxl.utils.column_index_from_string(c) for c in cols}
    lo, hi = min(idx.values()), max(idx.values())
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=r0, max_row=r1, min_col=lo, max_col=hi, values_only=True)):
        out.append((r0 + i, {c: row[k - lo] for c, k in idx.items()}))
    return out


def norm_lbl(s):
    return re.sub(r'\s', '', str(s))


class Model:
    def __init__(self, path):
        self.wb = openpyxl.load_workbook(path, data_only=True)
        wb = self.wb
        b = wb[SH_BASE]
        pv = lambda k: b[cell_of(P[k])].value
        self.year = int(N(pv('YEAR')))
        self.open = as_date(pv('OPEN'))
        self.st = max(self.open, dt.date(self.year, 1, 1))            # 本账从哪天起（建账日在上年就从 1/1）
        self.open_m = self.open.month if self.open.year == self.year else 1
        self.basis = pv('BASIS')
        self.cust0 = xtrim(xtext(pv('CUST')))
        self.lagd = N(pv('LAGD'))
        self.depto_in = N(pv('DEPTO'))
        # ⑩ 每月结账日期（ISNUMBER：日期或数字）
        self.close = {}
        for m in range(1, 13):
            v = b[f'{CL_DATE}{CL_R0 + m - 1}'].value
            self.close[m] = serial(v) if isnum(v) else None
        self.lock = max([m for m in range(1, 13) if self.close[m] is not None], default=0)
        # 账户（科目编码照 s_base：现金 1001；银行 1002xx（第几个，最多 6）；支付宝微信 1012xx（最多 4））
        self.acc = {}
        nb = npay = 0
        for r in range(AC_R0, AC_R1 + 1):
            nm, tp = b[f'{AC_NAME}{r}'].value, b[f'{AC_TYPE}{r}'].value
            if tp == '银行':
                nb += 1
            elif tp == '支付宝微信':
                npay += 1
            if nm in (None, ''):
                continue
            code = {'现金': '1001', '银行': f'1002{nb:02d}' if nb <= 6 else '',
                    '支付宝微信': f'1012{npay:02d}' if npay <= 4 else ''}.get(tp, '')
            self.acc[xtext(nm).casefold()] = dict(name=xtext(nm), code=code, open=N(b[f'{AC_OPEN}{r}'].value), k=r - AC_R0 + 1)
        self.acc_open_all = sum(N(b[f'{AC_OPEN}{r}'].value) for r in range(AC_R0, AC_R1 + 1))
        self.cats = {}
        for r in range(CT_R0, CT_R1 + 1):
            nm = b[f'{CT_NAME}{r}'].value
            if nm not in (None, '') and xtext(nm).casefold() not in self.cats:
                self.cats[xtext(nm).casefold()] = dict(opp=xtrim(xtext(b[f'{CT_OPP}{r}'].value)),
                                                      need=xtrim(xtext(b[f'{CT_NEED}{r}'].value)),
                                                      cf=xtrim(xtext(b[f'{CT_CF}{r}'].value)))
        self.fees = {}
        for r in range(FE_R0, FE_R1 + 1):
            nm = b[f'{FE_NAME}{r}'].value
            if nm not in (None, '') and xtext(nm).casefold() not in self.fees:
                cls = xtrim(xtext(b[f'{FE_CLS}{r}'].value))
                self.fees[xtext(nm).casefold()] = dict(cls=cls, code=FE_CLASSES.get(cls, ('',))[0])
        self.depts = {}
        for r in range(DP_R0, DP_R1 + 1):
            nm = b[f'{DP_NAME}{r}'].value
            if nm not in (None, '') and xtext(nm).casefold() not in self.depts:
                self.depts[xtext(nm).casefold()] = xtext(b[f'{DP_TYPE}{r}'].value)
        # 会计科目表：编码、上级、年初（用的 H 列）、末级
        coa = wb[SH_COA]
        self.coa = []
        for r, v in grab(coa, COA_R0, COA_R1, [COA_CODE, COA_UP, COA_OPEN]):
            self.coa.append(dict(r=r, code=xtrim(xtext(v[COA_CODE])), up=xtrim(xtext(v[COA_UP])), open=N(v[COA_OPEN])))
        ups = {x['up'].casefold() for x in self.coa if x['up']}
        self.coa_idx = {}
        for x in self.coa:
            x['leaf'] = x['code'] != '' and x['code'].casefold() not in ups
            if x['code'] and x['code'].casefold() not in self.coa_idx:
                self.coa_idx[x['code'].casefold()] = x
        # 往来单位
        u = wb[SH_UNIT]
        self.units = {}
        self.unit_by = {}
        for r, v in grab(u, UN_R0, UN_R1, [UN_NAME, UN_TYPE, UN_AP0, UN_AR0, UN_OR0, UN_OP0]):
            nm = v[UN_NAME]
            if nm in (None, ''):
                continue
            x = dict(name=xtext(nm), type=xtext(v[UN_TYPE]), ap0=N(v[UN_AP0]), ar0=N(v[UN_AR0]), or0=N(v[UN_OR0]), op0=N(v[UN_OP0]), row=r)
            self.units[xtext(nm)] = x
            self.unit_by.setdefault(xtext(nm).casefold(), x)
        # 款式
        s = wb[SH_STY]
        self.sty = []           # 按档案行
        self.sty_idx = {}
        for r, v in grab(s, ST_R0, ST_R1, [ST_CODE, ST_PRICE, ST_KM, ST_KL, ST_SDM, ST_SDL, ST_STAT]):
            key = xtrim(xtext(v[ST_CODE]))
            self.sty.append(dict(code=key, r=r, price=N(v[ST_PRICE]), status=xtrim(xtext(v[ST_STAT])),
                                 km=1 if v[ST_KM] in (None, '') else N(v[ST_KM]),
                                 kl=1 if v[ST_KL] in (None, '') else N(v[ST_KL]),
                                 sdm=N(v[ST_SDM]), sdl=N(v[ST_SDL])))
            if key and key.casefold() not in self.sty_idx:
                self.sty_idx[key.casefold()] = len(self.sty) - 1
            elif key:                      # 重复登记的编码：表里 _款式月 整行不算，这里也当空行
                self.sty[-1]['code'] = ''
        sdms = [x['sdm'] for x in self.sty if x['sdm'] > 0]
        sdls = [x['sdl'] for x in self.sty if x['sdl'] > 0]
        am, al_ = (sum(sdms) / len(sdms) if sdms else 1), (sum(sdls) / len(sdls) if sdls else 1)
        for x in self.sty:
            if not x['code']:
                x['wm'] = x['wl'] = 0
            elif self.basis == '标准成本':
                x['wm'] = x['sdm'] if x['sdm'] > 0 else x['km'] * am
                x['wl'] = x['sdl'] if x['sdl'] > 0 else x['kl'] * al_
            else:
                x['wm'], x['wl'] = x['km'], x['kl']
        self.je = []            # (月, 借, 贷, 金额, 单位, 来源, 日期)
        self.cost = []          # (类型 款/公, 款式行号 or None, 组件, 月, 金额)
        self._orders()          # 先算订单：双数决定停产款的 REL，REL 决定别的来源的成本记到款还是公共
        self._rel()
        self._journal()
        self._dn()
        self._out()
        self._wages()
        self._mj()
        self.lastm = min(12, max([0] + [j['m'] for j in self.jrows if j['ok']] + [x['cm'] for x in self.dn + self.ot]
                                 + [o['dm'] for o in self.od if o['ok']] + [w['m'] for w in self.wg] + [x['m'] for x in self.mjrows]))
        self._depr()
        self._engine()

    def sty_of(self, code):
        if blank(code):
            return None
        return self.sty_idx.get(tkey(code))

    def atype(self, si, cm):
        """"款"：登记的款式；否则 "公"（停产在成本引擎里转公共，见 _engine 的 R）"""
        return '款' if si is not None else '公'

    def month_of(self, d):
        d = as_date(d)
        return d.month if d and d.year == self.year else 0

    # ── 订单明细（成本双数 CQ、交货月 DM、金额 AMT、能记账 OK、订单号规范写法 NOK）
    def _orders(self):
        ws = self.wb[SH_ORD]
        cols = [OD_NO, OD_DATE, OD_CUST, OD_STY, OD_SPEC, OD_QTY, OD_NOTE, OD_DDATE, OD_DQTY, OD_PRICE]
        self.od = []
        for r, v in grab(ws, OD_R0, OD_R1, cols):
            if all(v[c] in (None, '') for c in cols):          # COUNTA(B:K)=0 → 空行
                continue
            g, j, k, dd = v[OD_QTY], v[OD_DQTY], v[OD_PRICE], v[OD_DDATE]
            si = self.sty_of(v[OD_STY])
            dq = (N(g) if blank(j) else N(j)) if isnum(dd) else 0.0
            ddate = as_date(dd) if isnum(dd) else None
            dm = 0
            if ddate and dq != 0 and ddate.year == self.year and int(serial(dd)) >= N(self.open):
                dm = ddate.month
            if isnum(k):
                price = N(k)
            elif not blank(k):
                price = 0.0
            else:
                price = self.sty[si]['price'] if si is not None else 0.0
            amt = d2(dq * price) if dm > 0 else 0.0
            bad = (blank(v[OD_NO]) or blank(v[OD_STY]) or (blank(g) and blank(j)) or (not blank(dd) and not isnum(dd))
                   or any(not blank(x) and not isnum(x) for x in (g, j, k)))
            ok = dm >= 1 and amt != 0 and not bad
            cq = (0.0 if '返修' in xtext(v[OD_NOTE]) else dq) if dq > 0 else 0.0
            cust = self.cust0 if blank(v[OD_CUST]) else xtrim(xtext(v[OD_CUST]))
            self.od.append(dict(r=r, nok=xtrim(xtext(v[OD_NO])), si=si, dq=dq, dm=dm, amt=amt if ok else 0.0, ok=ok, cust=cust,
                                oq=N(g), cq=cq, odate=serial(v[OD_DATE]) if isnum(v[OD_DATE]) else None,
                                ddate=serial(dd) if isnum(dd) else None))
            if ok:
                self.je.append((dm, '1122', '5001', amt, cust, '收入', None))
        n = len(self.sty)
        self.pairs = [[0.0] * 13 for _ in range(n)]
        self.rev = [[0.0] * 13 for _ in range(n)]
        for o in self.od:
            if o['ok'] and o['si'] is not None and self.sty[o['si']]['code']:
                self.pairs[o['si']][o['dm']] += o['cq']
                self.rev[o['si']][o['dm']] += o['amt']

    def _rel(self):
        """停产款：最后一个有双数的月份＋1（一双没交过＝1）；别的 13"""
        self.rel = []
        for i, x in enumerate(self.sty):
            if x['code'] and x['status'] == '停产':
                self.rel.append(max([m for m in range(1, 13) if self.pairs[i][m] > 0], default=0) + 1)
            else:
                self.rel.append(13)

    # ── 资金日记账
    def _journal(self):
        ws = self.wb[SH_CASH]
        cols = [J_DATE, J_ACC, J_CAT, J_UNIT, J_FEE, J_MEMO, J_IN, J_OUT, J_TOACC, J_STY, J_NO, J_NOTE]
        self.jrows = []
        for r, v in grab(ws, J_R0, J_R1, cols):
            seq = any(v[c] not in (None, '') for c in cols)
            hv, iv = txt2num(v[J_IN]), txt2num(v[J_OUT])
            net = d2((hv or 0.0) - (iv or 0.0))
            if not seq and net == 0:
                continue
            acc = self.acc.get(xtext(v[J_ACC]).casefold())
            unit = xtrim(xtext(v[J_UNIT]))
            un = self.unit_by.get(unit.casefold()) if unit else None
            utype = un['type'] if un else ''
            L = xtext(v[J_TOACC])
            cat = xtrim(xtext(v[J_CAT]))
            if not cat and net != 0:
                if L != '':
                    cat = XFER
                elif not blank(v[J_FEE]):
                    cat = '费用支出'
                else:
                    cat = {'材料供应商': '付材料款', '外发加工厂': '付加工费', '客户': '收货款'}.get(utype, '')
            row = dict(r=r, ok=False, net=net, acc=xtext(v[J_ACC]), cat=cat, to=L if cat == XFER else '', m=0, cf='')
            self.jrows.append(row)
            d = v[J_DATE]
            # ✗ 的条件（命中一条就进不了账）
            if blank(v[J_ACC]) or acc is None or not acc['code'] or not isnum(d) or as_date(d).year != self.year:
                continue
            if as_date(d) < self.open or (hv and iv) or hv is None or iv is None or hv < 0 or iv < 0 or net == 0:
                continue
            ct = self.cats.get(cat.casefold()) if cat else None
            if ct is None:
                continue
            need = ct['need']
            fee = self.fees.get(xtrim(xtext(v[J_FEE])).casefold()) if not blank(v[J_FEE]) else None
            to = self.acc.get(L.casefold()) if L else None
            if need == '单位' and (not unit or not utype):
                continue
            if need == '费用项目':
                if fee is None:
                    continue
                opp = fee['code']
            elif need == '对方账户':
                if to is None or to['k'] == acc['k']:
                    continue
                opp = to['code']
            else:
                opp = ct['opp']
            if not opp or opp.casefold() not in self.coa_idx:
                continue
            m = as_date(d).month
            row.update(ok=True, m=m)
            a = acc['code']
            dr, cr = (a, opp) if net > 0 else (opp, a)
            self.je.append((m, dr, cr, abs(net), unit, '日记账', as_date(d)))
            comp = COMP_OF.get(opp)
            if comp:
                si = self.sty_of(v[J_STY])
                self.cost.append((self.atype(si, m), si, comp, m, -net))
            if cat != XFER:
                if need == '费用项目':
                    row['cf'] = FE_CLASSES.get(fee['cls'], ('', '支付其他与经营有关的现金'))[1]
                else:
                    cf = ct['cf']
                    row['cf'] = CF_REPAY if (cf == CF_BORROW and net < 0) else (CF_BORROW if (cf == CF_REPAY and net > 0) else cf)

    # ── 送货单 / 外发
    def _cm(self, d, rd):
        """成本月（同 s_regs._bill_row）：0＝进不了账"""
        if not isnum(d):
            return 0, -1, -1
        dd = as_date(d)
        mm = lambda x: 0 if x.year < self.year else (13 if x.year > self.year else x.month)
        sm = mm(dd)
        rm = sm if blank(rd) else (mm(as_date(rd)) if isnum(rd) else -1)
        if sm == 13:
            return 0, sm, rm
        if dd < self.st:                                      # 建账前送的（含上年）
            if not isnum(rd) or as_date(rd) < self.st or rm == 13:
                return 0, sm, rm
            return max(rm, self.open_m), sm, rm
        cd = self.close[sm]
        if cd is not None:                                    # 送货月已结账：收单日期（空＝送货日期）晚于结账日期 → 收单月
            eff = int(serial(rd)) if isnum(rd) else int(serial(d))
            if eff > cd:
                return (0 if rm == 13 else rm), sm, rm
        return sm, sm, rm

    def _reg(self, sh, c_date, c_rdate, c_sup, c_qty, c_price, c_amt, c_sty, c_note, comp, r0, r1, code):
        ws = self.wb[sh]
        cols = [c_date, c_rdate, c_sup, c_qty, c_price, c_amt, c_sty]
        allc = [openpyxl.utils.get_column_letter(i) for i in range(openpyxl.utils.column_index_from_string(c_date),
                                                                   openpyxl.utils.column_index_from_string(c_note) + 1)]
        out = []
        for r, v in grab(ws, r0, r1, sorted(set(cols + allc), key=openpyxl.utils.column_index_from_string)):
            if all(v[c] in (None, '') for c in allc):
                continue
            d, rd, sup = v[c_date], v[c_rdate], v[c_sup]
            q_, p_, a_ = numv(v[c_qty]), numv(v[c_price]), v[c_amt]
            if isnum(a_):
                amt = d2(N(a_))
            elif blank(a_):
                amt = d2((q_ or 0) * (p_ or 0))
            else:
                amt = 0.0
            cm, sm, rm = self._cm(d, rd)
            supk = xtrim(xtext(sup))
            if (not isnum(d) or sm == 13 or blank(sup) or supk.casefold() not in self.unit_by
                    or (not blank(rd) and not isnum(rd)) or (isnum(rd) and int(serial(rd)) < int(serial(d)))
                    or cm < 1 or q_ is None or p_ is None or (not blank(a_) and not isnum(a_)) or amt == 0):
                continue
            si = self.sty_of(v[c_sty])
            self.cost.append((self.atype(si, cm), si, comp, cm, amt))
            self.je.append((cm, code, '2202', amt, supk, sh, as_date(d)))
            out.append(dict(r=r, cm=cm, sm=sm, rm=rm, amt=amt, sup=supk))
        return out

    def _dn(self):
        self.dn = self._reg(SH_DN, DN_DATE, DN_RDATE, DN_SUP, DN_QTY, DN_PRICE, DN_AMTIN, DN_STY, DN_NOTE, '材料', DN_R0, DN_R1,
                            COMP_CODE['材料'])

    def _out(self):
        self.ot = self._reg(SH_OUT, OT_DATE, OT_RDATE, OT_SUP, OT_QTY, OT_PRICE, OT_AMTIN, OT_STY, OT_NOTE, '外发', OT_R0, OT_R1,
                            COMP_CODE['外发'])

    # ── 工资
    def _wages(self):
        ws = self.wb[SH_WAGE]
        self.wg = []
        for r, v in grab(ws, WG_R0, WG_R1, [WG_MON, WG_NAME, WG_DEPT, WG_STY, WG_PQ, WG_PP, WG_BASE, WG_OT, WG_DED, WG_NOTE]):
            if all(v[c] in (None, '') for c in v):
                continue
            m = v[WG_MON]
            vals = [numv(v[c]) for c in (WG_PQ, WG_PP, WG_BASE, WG_OT, WG_DED)]
            if not isnum(m) or not 1 <= N(m) <= 12 or N(m) != int(N(m)) or N(m) < self.open_m:
                continue
            m = int(N(m))
            tp = self.depts.get(tkey(v[WG_DEPT])) if not blank(v[WG_DEPT]) else None
            if tp not in DP_TYPES or any(x is None for x in vals):
                continue
            pq, pp, base, ot, ded = vals
            piece = 0.0 if (blank(v[WG_PQ]) and blank(v[WG_PP])) else d2(pq * pp)
            amt = d2(piece + base + ot - ded)
            if amt == 0:
                continue
            comp, code, _ = DP_TYPES[tp]
            self.je.append((m, code, '2211', amt, '', '工资', None))
            if comp:
                si = self.sty_of(v[WG_STY])
                self.cost.append((self.atype(si, m), si, comp, m, amt))
            self.wg.append(dict(m=m, amt=amt, code=code))

    # ── 手工分录
    def _mj(self):
        ws = self.wb[SH_MJ]
        self.mjrows = []
        for r, v in grab(ws, MJ_R0, MJ_R1, [MJ_DATE, MJ_MEMO, MJ_DRIN, MJ_CRIN, MJ_AMTIN, MJ_UNIT, MJ_FEE, MJ_NOTE]):
            d = v[MJ_DATE]
            drc, crc = xtrim(xtext(v[MJ_DRIN])), xtrim(xtext(v[MJ_CRIN]))
            h = v[MJ_AMTIN]
            a = d2(N(h))
            if not isnum(d) or as_date(d).year != self.year or as_date(d) < self.open or not drc or not crc:
                continue
            xd, xc = self.coa_idx.get(drc.casefold()), self.coa_idx.get(crc.casefold())
            if xd is None or xc is None or not xd['leaf'] or not xc['leaf'] or drc == crc:
                continue
            if (h not in (None, '') and not isnum(h)) or a == 0 or drc in SYS_CODES or crc in SYS_CODES:
                continue
            if drc in COMP_OF and crc in COMP_OF:
                continue
            m = as_date(d).month
            self.je.append((m, drc, crc, a, xtrim(xtext(v[MJ_UNIT])), '手工', as_date(d)))
            self.mjrows.append(dict(r=r, m=m))
            if drc in COMP_OF:
                self.cost.append(('公', None, COMP_OF[drc], m, a))
            elif crc in COMP_OF:
                self.cost.append(('公', None, COMP_OF[crc], m, -a))

    # ── 固定资产折旧（提到 DEPTOX：填了按填的，空着＝最新月份 LASTM）
    def _depr(self):
        b = self.wb[SH_BASE]
        self.depto = min(12, int(self.depto_in)) if self.depto_in >= 1 else self.lastm
        self.dep = defaultdict(float)          # (m, use) → 折旧
        for r, v in grab(b, FA_R0, FA_R1, [FA_NAME, FA_DATE, FA_COST, FA_MON, FA_RES, FA_DEP0, FA_USE]):
            if v[FA_NAME] in (None, '') or not isnum(v[FA_DATE]):
                continue
            d0 = as_date(v[FA_DATE])
            cost, mon, res, dep0 = N(v[FA_COST]), N(v[FA_MON]), N(v[FA_RES]), N(v[FA_DEP0])
            if not cost or mon <= 0:
                continue
            dm = d2(cost * (1 - res) / mon)
            if dm == 0:
                continue
            start = d0.year * 12 + d0.month + 1
            t_open = self.open.year * 12 + self.open.month
            for m in range(1, self.depto + 1):
                t = self.year * 12 + m
                if t < t_open or t < start:
                    continue
                self.dep[(m, v[FA_USE])] += d2(max(0, min(dm, cost * (1 - res) - dep0 - dm * (t - max(start, t_open)))))

    # ── 成本引擎
    def _engine(self):
        n = len(self.sty)
        M = range(1, 13)
        direct = defaultdict(float)     # (si, comp, m)
        pub = defaultdict(float)        # (comp, m)
        for t, si, comp, m, a in self.cost:
            if t == '款':
                direct[(si, comp, m)] += a
            else:
                pub[(comp, m)] += a
        for m in M:
            pub[('制造', m)] += self.dep[(m, '车间')]
        self.pub = pub
        self.direct = direct
        self.direct_m = defaultdict(float)
        for (si, comp, m), a in direct.items():
            self.direct_m[(comp, m)] += a
        # 月末未交订单双数（只在本月有双数时算）：下单日期 < 下月 1 日（空也算）的订单数量 − 交货日期 < 下月 1 日的成本双数
        by_sty = defaultdict(list)
        for o in self.od:
            if o['si'] is not None:
                by_sty[o['si']].append(o)
        self.undel = [[0.0] * 13 for _ in range(n)]
        for si in range(n):
            if not self.sty[si]['code']:
                continue
            for m in M:
                if self.pairs[si][m] <= 0:
                    continue
                nxt = serial(dt.date(self.year + (m == 12), m % 12 + 1, 1))
                rows = by_sty[si]
                oq = sum(o['oq'] for o in rows) - sum(o['oq'] for o in rows if o['odate'] is not None and o['odate'] >= nxt)
                dq = sum(o['cq'] for o in rows if o['ddate'] is not None and o['ddate'] < nxt)
                self.undel[si][m] = max(0.0, oq - dq)
        # 直接转出：按交货进度
        T = defaultdict(float)
        for si in range(n):
            if not self.sty[si]['code']:
                continue
            for c in COMPS:
                done = cum = 0.0
                for m in M:
                    cum += direct[(si, c, m)]
                    p = self.pairs[si][m]
                    if p > 0:
                        T[(si, c, m)] = d2((cum - done) * p / (p + self.undel[si][m]))
                        done += T[(si, c, m)]
        self.T = T
        # 停产转公共：m ≥ REL 时，到本月累计的直接成本 − 1～m 月已转出 − 1～m−1 月已停转
        R = defaultdict(float)
        self.rel_in = defaultdict(float)
        for si in range(n):
            if not self.sty[si]['code'] or self.rel[si] > 12:
                continue
            for c in COMPS:
                cum = done = 0.0
                for m in M:
                    cum += direct[(si, c, m)]
                    done += T[(si, c, m)]
                    if m >= self.rel[si]:
                        R[(si, c, m)] = d2(cum - done)
                        done += R[(si, c, m)]
                        self.rel_in[(c, m)] += R[(si, c, m)]
        self.R = R
        wm = [x['wm'] for x in self.sty]
        wl = [x['wl'] for x in self.sty]
        self.base_m = {m: sum(self.pairs[i][m] * wm[i] for i in range(n)) for m in M}
        self.base_l = {m: sum(self.pairs[i][m] * wl[i] for i in range(n)) for m in M}
        a = self.wb[SH_ALLOC]
        self.wip_est = {m: N(a[f'{al_col(m)}{AL_ROWS["在制估计"]}'].value) for m in M}
        # 1 月「公共上月结转」＝年初在制（会计科目表 400101～400104 年初，H 列）
        self.wip0 = {c: d2(sum(x['open'] for x in self.coa if x['code'] == COMP_CODE[c])) for c in COMPS}
        self.alloc = defaultdict(float)
        self.keep = {}
        self.carry_in = {}
        self.carry = {}
        carry = dict(self.wip0)
        for m in M:
            avail = {c: d2(carry[c] + d2(pub[(c, m)]) + d2(self.rel_in[(c, m)])) for c in COMPS}
            tot = sum(avail.values())
            est = self.wip_est[m]
            for c in COMPS:
                base = self.base_m[m] if c == '材料' else self.base_l[m]
                keep = 0.0 if (est <= 0 or tot <= 0 or base <= 0) else d2(avail[c] * min(1, est / tot))
                self.carry_in[(c, m)] = carry[c]
                self.keep[(c, m)] = keep
                self.alloc[(c, m)] = d2(avail[c] - keep) if base > 0 else 0.0
                carry[c] = d2(avail[c] - self.alloc[(c, m)])
                self.carry[(c, m)] = carry[c]
        # 本X＝转X＋公共分到的（照表里：按累计加权双数逐行舍入，各款加起来正好＝公共本月分摊）
        self.C = defaultdict(float)
        self.cost_sm = [[0.0] * 13 for _ in range(n)]       # _款式月 成本
        self.unit = [[0.0] * 13 for _ in range(n)]          # _款式月 单双
        for m in M:
            for c in COMPS:
                base = self.base_m[m] if c == '材料' else self.base_l[m]
                w = wm if c == '材料' else wl
                pa = self.alloc[(c, m)]
                cw = 0.0
                for si in range(n):
                    p = self.pairs[si][m] if self.sty[si]['code'] else 0.0
                    prev = cw
                    cw += p * w[si]
                    if p <= 0:
                        continue
                    share = 0.0 if base <= 0 else d2(pa * cw / base) - d2(pa * prev / base)
                    self.C[(si, c, m)] = d2(T[(si, c, m)] + share)
            for si in range(n):
                p = self.pairs[si][m]
                self.cost_sm[si][m] = d2(sum(self.C[(si, c, m)] for c in COMPS))
                self.unit[si][m] = d4(self.cost_sm[si][m] / p) if p > 0 else 0.0
        self.direct_out = {(c, m): d2(sum(T[(si, c, m)] for si in range(n))) for c in COMPS for m in M}
        self.out_c = {(c, m): d2(self.direct_out[(c, m)] + self.alloc[(c, m)]) for c in COMPS for m in M}
        for m in M:
            self.je.append((m, COMP_CODE['制造'], MOH, d2(pub[('制造', m)] + self.direct_m[('制造', m)]), '', '制造转入', None))
            for c in COMPS:
                self.je.append((m, COGS_CODE[c], COMP_CODE[c], self.out_c[(c, m)], '', '成本结转', None))
            for use, code in FA_USES.items():
                if self.dep[(m, use)]:
                    self.je.append((m, code, '1602', d2(self.dep[(m, use)]), '', '折旧', None))
        # 订单行成本＝ROUND(成本双数 × 单双, 2)（能记账、款式登记了、成本双数>0 的行）
        for o in self.od:
            o['cost'] = d2(o['cq'] * self.unit[o['si']][o['dm']]) if (o['ok'] and o['si'] is not None and o['dm'] >= 1
                                                                      and o['cq'] > 0) else 0.0

    # ── 账
    def mov(self, prefix, m_from=1, m_to=12, unit=None, exact=False):
        hit = (lambda code: code == prefix) if exact else (lambda code: code.startswith(prefix))
        uk = None if unit is None else unit.casefold()
        ok_u = lambda x: uk is None or (x or '').casefold() == uk
        dr = sum(a for m, d, c, a, u, *_ in self.je if m_from <= m <= m_to and hit(d) and ok_u(u))
        cr = sum(a for m, d, c, a, u, *_ in self.je if m_from <= m <= m_to and hit(c) and ok_u(u))
        return dr, cr

    def pl(self, m1, m2):
        def cr_dr(p):
            d, c = self.mov(p, m1, m2)
            return c - d

        def dr_cr(p):
            d, c = self.mov(p, m1, m2)
            return d - c
        rev = cr_dr('5001') + cr_dr('5051')
        cost = dr_cr('5401') + dr_cr('5402')
        tax, sell, adm, fin = dr_cr('5403'), dr_cr('5601'), dr_cr('5602'), dr_cr('5603')
        oi, oe, it = cr_dr('5301'), dr_cr('5711'), dr_cr('5801')
        op = rev - cost - tax - sell - adm - fin
        return dict(收入=rev, 成本=cost, 毛利=rev - cost, 营业利润=op, 利润总额=op + oi - oe, 净利润=op + oi - oe - it,
                    材料=dr_cr('540101'), 外发=dr_cr('540102'), 人工=dr_cr('540103'), 制造=dr_cr('540104'))

    def bal(self, prefix, m, natural_debit=True, open_=0.0, unit=None):
        d, c = self.mov(prefix, 1, m, unit)
        return open_ + (d - c if natural_debit else c - d)

    def unit_now(self, x):
        """往来单位 M 列「现在余额」：(应付 2202＋其他应付 2241) − (应收 1122＋其他应收 1221)，我们欠他为正；客户取反"""
        v = x['ap0'] + x['op0'] - x['ar0'] - x['or0']
        for code in ('2202', '2241', '1122', '1221'):
            d, c = self.mov(code, 1, 12, x['name'], exact=True)
            v += c - d
        return -v if x['type'] == '客户' else v


# ─────────────────────────── 跟表里比 ───────────────────────────
def find_label(ws, test, max_row=120, max_col=30):
    """找第一个文字满足 test(去掉空白的文字) 的格子 → (行, 列号)"""
    for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
        for c in row:
            if isinstance(c.value, str) and test(norm_lbl(c.value)):
                return c.row, c.column
    return None


def month_cols(ws, max_row=8):
    mcol = {}
    for row in ws.iter_rows(min_row=1, max_row=max_row):
        for cc in row:
            if isinstance(cc.value, str) and re.fullmatch(r'(1[0-2]|[1-9])月', cc.value.strip()):
                mcol.setdefault(int(cc.value.strip()[:-1]), cc.column_letter)
    return mcol


def compare(path):
    md = Model(path)
    wb = md.wb
    bad, n_ok = [], 0

    def chk(what, got, exp, tol=TOL):
        nonlocal n_ok
        g = got if isinstance(got, (int, float)) else (0 if got in (None, '') else got)
        if not isinstance(g, (int, float)) or abs(g - exp) > tol:
            bad.append(f'{what}: 表里 {got!r}，复算 {exp:.2f}')
        else:
            n_ok += 1

    def chk_txt(what, got, ok_, exp_desc):
        nonlocal n_ok
        if ok_:
            n_ok += 1
        else:
            bad.append(f'{what}: 表里 {got!r}，应该 {exp_desc}')

    # ── 基础资料：自动参数
    b = wb[SH_BASE]
    chk('基础资料 最新月份 LASTM', b[cell_of(P['LASTM'])].value, md.lastm, 0)
    chk('基础资料 已结账到几月 LOCK（显示用）', b[cell_of(P['LOCK'])].value, md.lock, 0)
    chk('基础资料 折旧实际提到 DEPTOX', b[cell_of(P['DEPTOX'])].value, md.depto, 0)
    # ── _款式月
    sm = wb[SH_SM]
    for i, x in enumerate(md.sty):
        r = SM_R0 + i
        if not x['code']:
            continue
        for m in range(1, 13):
            chk(f'_款式月 {x["code"]} {m}月 双数', sm[f'{sm_col("双数", m)}{r}'].value, md.pairs[i][m])
            chk(f'_款式月 {x["code"]} {m}月 收入', sm[f'{sm_col("收入", m)}{r}'].value, md.rev[i][m])
            chk(f'_款式月 {x["code"]} {m}月 未交', sm[f'{sm_col("未交", m)}{r}'].value, md.undel[i][m])
            for c in COMPS:
                chk(f'_款式月 {x["code"]} {m}月 转{c}', sm[f'{sm_col("转" + c, m)}{r}'].value, md.T[(i, c, m)])
                chk(f'_款式月 {x["code"]} {m}月 停{c}', sm[f'{sm_col("停" + c, m)}{r}'].value, md.R[(i, c, m)])
                chk(f'_款式月 {x["code"]} {m}月 本{c}', sm[f'{sm_col("本" + c, m)}{r}'].value, md.C[(i, c, m)], 0.02)
            chk(f'_款式月 {x["code"]} {m}月 成本', sm[f'{sm_col("成本", m)}{r}'].value, md.cost_sm[i][m], 0.03)
    # ── 款式档案：已交双数、停产 REL、在制（挂着没转的直接成本）
    st = wb[SH_STY]
    for i, x in enumerate(md.sty):
        if not x['code']:
            continue
        r = x['r']
        chk(f'款式档案 {x["code"]} 已交双数', st[f'{ST_DQ}{r}'].value, sum(md.pairs[i][1:]))
        chk(f'款式档案 {x["code"]} REL（停产后转公共的起始月）', st[f'{ST_REL}{r}'].value, md.rel[i], 0)
        wip = (sum(md.direct[(i, c, m)] for c in COMPS for m in range(1, 13)) - sum(md.T[(i, c, m)] for c in COMPS for m in range(1, 13))
               - sum(md.R[(i, c, m)] for c in COMPS for m in range(1, 13)))
        chk(f'款式档案 {x["code"]} 在制（{ST_WIP} 列 挂着没转）', st[f'{ST_WIP}{r}'].value, d2(wip))
    # ── 成本分摊表
    al = wb[SH_ALLOC]
    for m in range(1, 13):
        chk(f'成本分摊表 {m}月 双数', al[f'{al_col(m)}{AL_ROWS["双数"]}'].value, sum(md.pairs[i][m] for i in range(len(md.sty))))
        for c in COMPS:
            g = lambda it: al[f'{al_col(m)}{AL_CROWS[c][it]}'].value
            chk(f'成本分摊表 {c} {m}月 本月发生', g('本月发生'), md.pub[(c, m)] + md.direct_m[(c, m)])
            chk(f'成本分摊表 {c} {m}月 公共上月结转', g('公共上月结转'), md.carry_in[(c, m)])
            chk(f'成本分摊表 {c} {m}月 停产转入', g('停产转入'), md.rel_in[(c, m)])
            chk(f'成本分摊表 {c} {m}月 公共月末留在制', g('公共月末留在制'), md.keep[(c, m)])
            chk(f'成本分摊表 {c} {m}月 公共本月分摊', g('公共本月分摊'), md.alloc[(c, m)])
            chk(f'成本分摊表 {c} {m}月 直接本月转出', g('直接本月转出'), md.direct_out[(c, m)])
            chk(f'成本分摊表 {c} {m}月 本月转出合计', g('本月转出合计'), md.out_c[(c, m)])
    # ── 订单明细：每行交货金额、成本；订单汇总：按订单号规范写法分组
    od = wb[SH_ORD]
    for o in md.od:
        chk(f'订单明细 第{o["r"]}行 交货金额(隐藏AMT)', od[f'{OD_AMT}{o["r"]}'].value, o['amt'])
        chk(f'订单明细 第{o["r"]}行 成本双数(隐藏CQ)', od[f'{OD_CQ}{o["r"]}'].value, o['cq'], 0)
        if o['ok']:
            chk(f'订单明细 第{o["r"]}行 成本', od[f'{OD_COSTV}{o["r"]}'].value, o['cost'])
    for c in od[3]:                                   # 第 3 行「已交双数」＝SUM(CQ)
        if isinstance(c.value, str) and norm_lbl(c.value) == '已交双数':
            got = next((od.cell(3, k).value for k in range(c.column + 1, c.column + 4) if od.cell(3, k).value is not None), None)
            chk('订单明细 第3行 已交双数（成本双数合计）', got, sum(o['cq'] for o in md.od), 0)
            break
    else:
        bad.append('订单明细：第 3 行找不到「已交双数」')
    tot_line = sum(od[f'{OD_COSTV}{o["r"]}'].value for o in md.od if isinstance(od[f'{OD_COSTV}{o["r"]}'].value, (int, float)))
    chk('订单明细 成本列合计 vs 复算款式成本合计', tot_line, sum(md.cost_sm[i][m] for i in range(len(md.sty)) for m in range(1, 13)),
        max(1, 0.005 * sum(1 for o in md.od if o['cost'])))
    groups = {}
    for o in md.od:
        if not o['nok']:
            continue
        g = groups.setdefault(o['nok'].casefold(), dict(no=o['nok'], oq=0.0, cq=0.0, amt=0.0, cost=0.0))
        g['oq'] += o['oq']
        g['cq'] += o['cq']
        g['amt'] += o['amt']
        g['cost'] += o['cost']
    os_ = wb[SH_OSUM]
    hit = find_label(os_, lambda t: t == '订单号', 10)
    if hit is None:
        bad.append('订单汇总：找不到「订单号」表头')
    else:
        hr = hit[0]
        heads = {norm_lbl(c.value): c.column_letter for c in os_[hr] if isinstance(c.value, str)}
        need = {'订单号': '订单号', '订单双数': '订单双数', '已交双数': '已交双数', '交货收入': '交货收入', '成本': '成本'}
        miss = [k for k in need if k not in heads]
        if miss:
            bad.append(f'订单汇总：表头找不到 {miss}')
        else:
            seen = []
            for r in range(hr + 1, hr + 2 + len(groups) + 5):
                no = os_[f'{heads["订单号"]}{r}'].value
                if blank(no):
                    continue
                k = tkey(no)
                seen.append(k)
                g = groups.get(k)
                if g is None:
                    bad.append(f'订单汇总 第{r}行 订单号 {no!r}：订单明细里没有这张单（或重复列出）')
                    continue
                for lab, key in (('订单双数', 'oq'), ('已交双数', 'cq'), ('交货收入', 'amt'), ('成本', 'cost')):
                    chk(f'订单汇总 {g["no"]} {lab}', os_[f'{heads[lab]}{r}'].value, d2(g[key]), 0.011 if key != 'cost' else 0.02)
            chk_txt('订单汇总 订单张数（按订单号规范写法，不重复）', len(seen), sorted(seen) == sorted(groups) and len(seen) == len(set(seen)),
                    f'{len(groups)} 张')
    # ── 资金日记账：本账户余额 J、总余额 K（逐行即时，✗ 行也算进去）
    cash = wb[SH_CASH]
    run_acc = defaultdict(float)
    run_tot = 0.0
    for j in md.jrows:
        run_acc[j['acc'].casefold()] += j['net']
        if j['to']:
            run_acc[j['to'].casefold()] -= j['net']
        if j['cat'] != XFER:
            run_tot += j['net']
        if j['net'] == 0:
            continue
        r = j['r']
        if j['acc'] != '':
            a = md.acc.get(j['acc'].casefold())
            chk(f'资金日记账 第{r}行 本账户余额', cash[f'{J_BAL}{r}'].value, d2((a['open'] if a else 0) + run_acc[j['acc'].casefold()]))
        chk(f'资金日记账 第{r}行 总余额', cash[f'{J_TBAL}{r}'].value, d2(md.acc_open_all + run_tot))
    # ── 账户余额
    for r in range(AC_R0, AC_R1 + 1):
        nm = b[f'{AC_NAME}{r}'].value
        if nm in (None, ''):
            continue
        a = md.acc[xtext(nm).casefold()]
        k = xtext(nm).casefold()
        exp = a['open'] + sum(j['net'] for j in md.jrows if j['acc'].casefold() == k) - sum(j['net'] for j in md.jrows if j['to'].casefold() == k)
        chk(f'基础资料 {nm} 当前余额', b[f'{AC_NOW}{r}'].value, exp)
    # ── 往来单位现在余额（全部已记的：(2202＋2241)−(1122＋1221)，客户取反）；首页往来四个数
    u = wb[SH_UNIT]
    four = dict(应付=0.0, 预付=0.0, 应收=0.0, 预收=0.0)
    for nm, x in md.units.items():
        v = md.unit_now(x)
        chk(f'往来单位 {nm} 现在余额', u[f'M{x["row"]}'].value, v)
        v = d2(v)
        if x['type'] in SUPS:
            four['应付' if v > 0 else '预付'] += abs(v)
        elif x['type'] == '客户':
            four['应收' if v > 0 else '预收'] += abs(v)
    hm = wb[SH_HOME]
    tests = {'应付': lambda t: '（应付）' in t and '供应商' in t, '预付': lambda t: '预付给供应商' in t,
             '应收': lambda t: '（应收）' in t, '预收': lambda t: '（预收）' in t}
    for k, tst in tests.items():
        hit = find_label(hm, tst, 200, 20)
        if hit is None:
            bad.append(f'首页：找不到往来「{k}」那一格')
        else:
            chk(f'首页 往来 {k}', hm.cell(hit[0] + 1, hit[1]).value, four[k])
    # ── 利润表：按标签找行，1～12 月
    is_ = wb[SH_IS]
    lab = {}
    for row in is_.iter_rows(min_row=1, max_row=80):
        for cc in row[:3]:
            if isinstance(cc.value, str):
                t = re.sub(r'[\s一二三四五六七八九十、：:（）()减加其中]', '', cc.value)
                lab.setdefault(t, cc.row)
    mcol = month_cols(is_)
    want = {'营业收入': '收入', '营业成本': '成本', '营业利润': '营业利润', '利润总额': '利润总额', '净利润': '净利润'}
    if len(mcol) == 12:
        for k, key in want.items():
            if k in lab:
                for m in range(1, 13):
                    chk(f'利润表 {k} {m}月', is_[f'{mcol[m]}{lab[k]}'].value, md.pl(m, m)[key], 0.02)
            else:
                bad.append(f'利润表：找不到「{k}」这一行')
    else:
        bad.append(f'利润表：找不到 1～12 月的表头（找到 {sorted(mcol)}）')
    # ── 现金流量表：每个项目 × 月、期末现金
    cf = wb[SH_CF]
    mcol = month_cols(cf)
    rows = {}
    for row in cf.iter_rows(min_row=1, max_row=60, max_col=1):
        c = row[0]
        if isinstance(c.value, str):
            rows.setdefault(norm_lbl(c.value), c.row)
    if len(mcol) == 12:
        for nm, _g in CF_ITEMS:
            if nm not in rows:
                bad.append(f'现金流量表：找不到「{nm}」这一行')
                continue
            for m in range(1, 13):
                exp = sum(j['net'] for j in md.jrows if j['ok'] and j['m'] == m and j['cf'] == nm)
                chk(f'现金流量表 {nm} {m}月', cf[f'{mcol[m]}{rows[nm]}'].value, exp)
        er = rows.get('五、期末现金余额')
        if er:
            for m in range(1, 13):
                exp = md.acc_open_all + sum(j['net'] for j in md.jrows if j['ok'] and 1 <= j['m'] <= m and j['cat'] != XFER)
                chk(f'现金流量表 期末现金余额 {m}月', cf[f'{mcol[m]}{er}'].value, exp)
        else:
            bad.append('现金流量表：找不到「五、期末现金余额」')
    else:
        bad.append(f'现金流量表：找不到 1～12 月的表头（找到 {sorted(mcol)}）')
    # ── 科目余额表：所选月份（C3 空＝最新月份）每个科目本月借贷、本年累计；最下面核对格 √
    tb = wb[SH_TB]
    c3 = tb['C3'].value
    mx = int(N(c3)) if isnum(c3) and 1 <= N(c3) <= 12 else md.lastm
    if mx >= 1:
        for r in range(COA_R0, COA_R1 + 1):
            code = xtrim(xtext(tb[f'A{r}'].value))
            if not code:
                continue
            d, c = md.mov(code, mx, mx)
            dc, cc = md.mov(code, 1, mx)
            chk(f'科目余额表 {code} {mx}月 借方', tb[f'F{r}'].value, d)
            chk(f'科目余额表 {code} {mx}月 贷方', tb[f'G{r}'].value, c)
            chk(f'科目余额表 {code} 本年借方累计到{mx}月', tb[f'I{r}'].value, dc)
            chk(f'科目余额表 {code} 本年贷方累计到{mx}月', tb[f'J{r}'].value, cc)
    hit = None
    for r in range(COA_R1 + 1, COA_R1 + 12):
        if norm_lbl(tb[f'A{r}'].value or '') == '核对':
            hit = r
            break
    if hit is None:
        bad.append('科目余额表：最下面找不到「核对」行')
    else:
        v = tb[f'C{hit}'].value
        chk_txt(f'科目余额表 核对格（C{hit}）', v, isinstance(v, str) and v.startswith('√'), '√ 开头')
    # ── 资产负债表：实际月份 BS_MX、平衡
    bs = wb[SH_BS]
    bm = bs[cell_of(BS_M)].value
    chk('资产负债表 实际月份 BS_MX', bs[cell_of(BS_MX)].value, int(N(bm)) if isnum(bm) and 1 <= N(bm) <= 12 else md.lastm, 0)
    chk('资产负债表 BS_CHECK（应为 0）', bs[cell_of(BS_CHECK)].value, 0)
    # ── 生产成本余额 = 累计发生 − 累计转出（12 月末，都不含年初）
    wip = sum(md.pub[(c, m)] + md.direct_m[(c, m)] for c in COMPS for m in range(1, 13)) - sum(md.out_c.values())
    d, c = md.mov('4001')
    chk('生产成本 4001 本年发生净额（借−贷）', d - c, wip, 0.05)
    print(f'核对 {n_ok + len(bad)} 项：对上 {n_ok}，不对 {len(bad)}')
    for x in bad[:80]:
        print('  ✗', x)
    return md, bad


if __name__ == '__main__':
    compare(sys.argv[1])
