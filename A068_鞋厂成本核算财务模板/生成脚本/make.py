# -*- coding: utf-8 -*-
"""把各张表拼成一本：首页 → 录入（蓝）→ 成本查看（绿）→ 往来资金查看（绿）→ 报表（深红）→ 数据校验 → 档案（灰）；隐藏的取数表放最后。
   先按顺序把所有表建好（空表），再让各模块往里写——模块之间互相引用的表一定存在。"""
import importlib
import re
import openpyxl
from openpyxl.styles import Protection
from openpyxl.worksheet.protection import SheetProtection
from common import home_link, print_setup, CL
from layout import *
import demo

GROUPS = [
    ('首页', C_HOME, [SH_HOME]),
    ('录入', C_IN, [SH_CASH, SH_DN, SH_OUT, SH_ORD, SH_WAGE, SH_MJ]),
    ('查看', C_VIEW, [SH_SPL, SH_OSUM, SH_OQ, SH_ALLOC, SH_MSUM, SH_LATE,
                     SH_APS, SH_SST, SH_CST, SH_ACCB, SH_FEE, SH_WSUM]),
    ('报表', C_RPT, [SH_IS, SH_BS, SH_CF, SH_TB, SH_GL, SH_JE]),
    ('校验', C_CHK, [SH_CHK]),
    ('档案', C_ARC, [SH_BASE, SH_UNIT, SH_STY, SH_MAT, SH_COA]),
]
HIDDEN = [SH_SM, SH_AUX]
MODULES = ['s_base', 's_cash', 's_regs', 's_cost', 's_ledger', 's_views', 's_home']
KEEP_OPEN = {'FFFFFF00', 'FFFFF7E0', 'FFEAF1FB'}      # 黄色选择格、淡黄手填格、淡蓝粘贴格：保护时不锁


def _balanced(f):
    depth, q_ = 0, False
    for ch in f:
        if ch == '"':
            q_ = not q_
        elif not q_:
            depth += (ch == '(') - (ch == ')')
            if depth < 0:
                return False
    return depth == 0 and not q_


def _fit_widths(wb):
    """日期、金额列放宽，免得 Excel/WPS 里显示成 ####（微软雅黑的数字比默认字体宽）"""
    for ws in wb.worksheets:
        need = {}
        anchors = {(m.min_row, m.min_col) for m in ws.merged_cells.ranges if m.max_col > m.min_col}
        for row in ws.iter_rows():
            for c in row:
                if (c.row, c.column) in anchors:
                    continue
                fmt = c.number_format or ''
                sz = (c.font.sz or 10) if c.font is not None else 10
                if c.value is None or ws.column_dimensions[c.column_letter].hidden:
                    continue
                w = 0
                if 'yy' in fmt:
                    w = 12.5 if 'm/d' not in fmt or 'yyyy' in fmt else 8
                elif '#,##0.00' in fmt:
                    w = 13.5
                elif '#,##0' in fmt:
                    w = 11
                if w:
                    w = w * max(1.0, sz / 10.0) + (1.5 if (c.font is not None and c.font.b) else 0)
                    need[c.column_letter] = max(need.get(c.column_letter, 0), w)
        for col, w in need.items():
            cd = ws.column_dimensions[col]
            if (cd.width or 8.43) < w:
                cd.width = w


BANNED = re.compile(r'(?<![A-Z0-9_.])(XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|LET|LAMBDA|MAXIFS|MINIFS|IFS|SWITCH|'
                    r'TEXTJOIN|CONCAT|INDIRECT|OFFSET|TODAY|NOW|RAND|RANDBETWEEN)\(')


def _check_formulas(wb):
    """只用 Excel 2010 / 老版 WPS 都有的函数；不用易变函数；括号、引号要配对"""
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if isinstance(v, str) and v.startswith('='):
                    outside = ''.join(v.split('"')[0::2]).upper()
                    hit = BANNED.findall(outside) + (['_xlfn'] if '_XLFN.' in outside else [])
                    if hit or not _balanced(v):
                        bad.append(f'{ws.title}!{c.coordinate}: {hit} {v[:120]}')
    assert not bad, '\n'.join(bad[:30])


def build(only=None, scenario=None):
    """only：只跑这几个模块（调试用）；scenario(ctx)：改演示数据（场景测试用）"""
    ctx = demo.build_ctx()
    if scenario:
        scenario(ctx)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for _g, color, names in GROUPS:
        for n in names:
            ws = wb.create_sheet(n)
            ws.sheet_properties.tabColor = color[2:]
    for n in HIDDEN:
        wb.create_sheet(n)
    for m in MODULES:
        if only and m not in only:
            continue
        try:
            mod = importlib.import_module(m)
        except ModuleNotFoundError as e:
            if e.name == m:
                print(f'   （模块 {m} 还没写，跳过）')
                continue
            raise
        mod.build(wb, ctx)
    for n in HIDDEN:
        wb[n].sheet_state = 'hidden'
    for grp, color, names in GROUPS:
        for n in names:
            ws = wb[n]
            ws.sheet_properties.tabColor = color[2:]
            if grp in ('录入', '档案'):
                end = max((mg.max_col for mg in ws.merged_cells.ranges if mg.min_row == 1 and mg.min_col == 1), default=ws.max_column)
                home_link(ws, f'{CL(end + 1)}1')
                ws.column_dimensions[CL(end + 1)].width = max(ws.column_dimensions[CL(end + 1)].width or 0, 10)
            if grp in ('查看', '报表', '校验', '首页'):
                for row in ws.iter_rows():
                    for c in row:
                        if c.fill is not None and c.fill.fgColor is not None and c.fill.fgColor.rgb in KEEP_OPEN:
                            c.protection = Protection(locked=False)
                ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                                formatCells=False, selectLockedCells=False, selectUnlockedCells=False)
            if not ws.print_title_rows and grp != '首页':
                hdr = {'录入': '4:4' if n != SH_CASH else f'{J_HDR}:{J_HDR}', '档案': '4:4'}.get(grp)
                print_setup(ws, hdr, landscape=True)
    _check_formulas(wb)
    _fit_widths(wb)
    for ws in wb.worksheets:
        ws.sheet_view.tabSelected = False
    wb.active = 0
    wb.worksheets[0].sheet_view.tabSelected = True
    wb.calculation.fullCalcOnLoad = False
    wb.calculation.calcId = 191029
    return wb
