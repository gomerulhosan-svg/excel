# -*- coding: utf-8 -*-
"""成本核心（第二阶段）：【_款式月】（隐藏取数表）【成本分摊表】【款式成本利润】【订单汇总】【订单查询】。

核算办法（开发说明「模块 s_cost」）：
  · 款式 × 月：直接记到款式的成本（CKEY＝"款|款式|组件"）本月有交货才转出（到本月累计 − 以前已转出），没交货挂着；
    公共成本（CKEY＝"公||组件"，制造另加车间折旧）本月有交货就按「交货双数 × 权重」全部分掉（减手填的「在制估计」），
    一双都没交的月份整月结转下月。
  · 分公共成本时每个款的份额用「累计舍入」：第 i 个款分到 ROUND(公共×前 i 个款累计加权双数÷加权双数,2) − ROUND(公共×前 i−1 个…,2)，
    每个款都是 2 位小数、跟精确值差不到 1 分，而且各款加起来正好等于「公共本月分摊」——成本分摊表「核对」行 12 个月都是 0。
  · 本表内部的格子（帮手列、选择格）自己定；被别的表引用的格子按 layout.py（SM_*、AL_*、OQ_SEL）。
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from common import *
from layout import *          # 注意：layout 的 C_IN 等颜色覆盖 common 的同名常量

H_VIEW = 'FF548235'                         # 查看表表头：绿
H_DARK = 'FF375623'
KPI_FILL = fill('FFD9E1F2')
F_GOOD = Font(name=YH, sz=10, bold=True, color='FF00B050')
F_WAIT = Font(name=YH, sz=10, color='FF7F7F7F')
F_REDN = Font(name=YH, sz=10, color='FFC00000')
FILL_YEL = fill('FFFFEB9C')
CHKF = '"✗ 差 "#,##0.00;"✗ 差 "-#,##0.00;"√"'      # 核对格：0 显示 √
WT = '0.00;-0.00;"-"'
MONTHF = '0"月"'

PERIOD = P['PERIOD']
SMQ = q(SH_SM)


def _w(comp):
    """组件 → (_款式月 权重列, 成本分摊表 加权双数行名)"""
    return (SM_WM, '加权材料') if comp == '材料' else (SM_WL, '加权工费')


def _od(col, r):
    """【订单明细】某列第 r 行（单格，绝对列）"""
    return f'{q(SH_ORD)}!${col}{r}'


def _odg(col, r):
    """【订单明细】某列「表头行 ～ 第 r 行」（往下拖时变长）"""
    return f'{q(SH_ORD)}!${col}${OD_HDR}:${col}{r}'


def _lbl(ws, coord, text, merge=None, fill_=KPI_FILL):
    if merge:
        ws.merge_cells(merge)
    put(ws, coord, text, F_KPI_L, fill_, align=ACW)


def _val(ws, coord, f, fmt=MONEY, merge=None, font=F_KPI_V):
    if merge:
        ws.merge_cells(merge)
    put(ws, coord, f, font, KPI_FILL, fmt, AC)


def _chk_cf(ws, rg, first):
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="✗"'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="⚠"'], fill=FILL_YEL))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="√"'], font=F_GOOD))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT({first},1)="⏳"'], font=F_WAIT))


def _neg_red(ws, rg, first):
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'AND(ISNUMBER({first}),{first}<0)'], font=F_REDN))


def _help_cols(ws, hdr_row, labels, r0, r1):
    """隐藏帮手列：表头小灰字、整列小灰字、隐藏"""
    for c, t in labels:
        ws[f'{c}{hdr_row}'] = t
        ws[f'{c}{hdr_row}'].font = F_HELP
        for r in range(r0, r1 + 1):
            ws[f'{c}{r}'].font = F_HELP
    hide(ws, *[c for c, _ in labels])


# ═══════════════════════════ _款式月（隐藏） ═══════════════════════════
def build_sm(wb, ctx):
    ws = wb[SH_SM]
    r0, r1 = SM_R0, SM_R1
    ws['A1'] = ('_款式月（隐藏取数表，全自动）：第 r 行＝【款式档案】第 r 行。每块 12 列＝1～12 月（第 3 行是月份数字）。'
                '双数/收入来自【订单明细】；转X＝直接成本本月转出；本X＝转X＋分到的公共成本；成本＝四块合计；单双＝成本÷双数。')
    ws['A1'].font = F_TXTB
    ws['A2'] = '别的表按 layout.py 取数，不要改动、不要插行删行。D 列＝款式序号（【订单明细】SROW 就是它）。'
    ws['A2'].font = F_NOTE
    for c, t in ((SM_CODE, '款式编码'), (SM_WM, '材料权重'), (SM_WL, '工费权重'), ('D', '款式序号')):
        put(ws, f'{c}{SM_HDR}', t, F_HDR, fill(H_DARK), align=ACW)
    ws.column_dimensions[SM_CODE].width = 12
    for i, blk in enumerate(SM_BLOCKS):
        color = 'FF548235' if i % 2 == 0 else 'FF2F75B5'
        for m in range(1, 13):
            c = sm_col(blk, m)
            put(ws, f'{c}{SM_MROW}', m, F_AUTOB, align=AC)
            put(ws, f'{c}{SM_HDR}', f'{blk}{m}', F_HDR, fill(color), align=AC)
            ws.column_dimensions[c].width = 9
    od_dq, od_srow, od_dm, od_ok, od_amt = odr(OD_DQ), odr(OD_SROW), odr(OD_DM), odr(OD_OK), odr(OD_AMT)
    for r in range(r0, r1 + 1):
        A, D = f'${SM_CODE}{r}', f'$D{r}'
        key = f'{q(SH_STY)}!${ST_KEY}{r}'
        ws[f'{SM_CODE}{r}'] = f'=IF({key}="","",{key}&"")'
        ws[f'{SM_WM}{r}'] = f'=IF({A}="",0,N({q(SH_STY)}!${ST_WM}{r}))'
        ws[f'{SM_WL}{r}'] = f'=IF({A}="",0,N({q(SH_STY)}!${ST_WL}{r}))'
        ws[f'D{r}'] = f'=ROW()-{r0 - 1}'
        for m in range(1, 13):
            cq, cr = sm_col('双数', m), sm_col('收入', m)
            ws[f'{cq}{r}'] = (f'=IF({A}="",0,SUMIFS({od_dq},{od_srow},{D},{od_dm},{cq}${SM_MROW},{od_dq},">0",{od_ok},1))')
            ws[f'{cr}{r}'] = f'=IF({A}="",0,ROUND(SUMIFS({od_amt},{od_srow},{D},{od_dm},{cr}${SM_MROW},{od_ok},1),2))'
            dq = f'{cq}{r}'
            # 转X：本月有交货 → 到本月累计的直接成本 − 1～m−1 月已转出；没交货 0（IF 短路：双数 0 时不算 cost_sum）
            for comp in COMPS:
                ct = sm_col(f'转{comp}', m)
                cs = cost_sum(f'"款|"&{esc(A)}&"|{comp}"', f'"<="&{ct}${SM_MROW}', comp)
                prior = '' if m == 1 else f'-SUM(${sm_col(f"转{comp}", 1)}{r}:{sm_col(f"转{comp}", m - 1)}{r})'
                ws[f'{ct}{r}'] = f'=IF({A}="",0,IF({dq}<=0,0,ROUND({cs}{prior},2)))'
            # 本X：转X ＋ 公共本月分摊 × 本款加权双数 ÷ 全部加权双数（累计舍入，各款之和正好＝公共本月分摊）
            for comp in COMPS:
                cb, ct = sm_col(f'本{comp}', m), sm_col(f'转{comp}', m)
                wcol, wrow = _w(comp)
                pub, base, w = al(comp, '公共本月分摊', m), alr(wrow, m), f'${wcol}{r}'
                cw = f'SUMPRODUCT(${cq}${r0}:{cq}{r},${wcol}${r0}:${wcol}{r})'
                share = f'ROUND({pub}*{cw}/{base},2)-ROUND({pub}*({cw}-{dq}*{w})/{base},2)'
                ws[f'{cb}{r}'] = f'=IF({A}="",0,IF({dq}<=0,0,ROUND({ct}{r}+IF({base}<=0,0,{share}),2)))'
            cc, cu = sm_col('成本', m), sm_col('单双', m)
            ws[f'{cc}{r}'] = '=ROUND(' + '+'.join(f'{sm_col("本" + comp, m)}{r}' for comp in COMPS) + ',2)'
            ws[f'{cu}{r}'] = f'=IF({dq}<=0,0,ROUND({cc}{r}/{dq},4))'
    ws.freeze_panes = f'{SM_B0}{r0}'


# ═══════════════════════════ 成本分摊表 ═══════════════════════════
_ITEM_LBL = {'本月发生': '{c} 本月发生', '其中直接到款式': '　其中：直接记到款式', '其中公共': '　　　　公共（要分摊）',
             '公共上月结转': '公共：上月结转来', '公共本月可分摊': '公共：本月可分摊', '公共月末留在制': '公共：月末留在制',
             '公共本月分摊': '公共：本月分到各款', '公共结转下月': '公共：结转下月', '直接本月转出': '直接：本月转出（交货的款）',
             '本月转出合计': '{c} 本月转出合计'}
_COMP_NAME = {'材料': '材料', '外发': '外发加工', '人工': '直接人工', '制造': '制造费用'}
_SRC_NOTE = {'材料': '送货单＋现付材料（日记账）＋手工分录，按成本月份',
             '外发': '外发加工登记＋现付加工费（日记账）＋手工分录',
             '人工': '工资登记「直接生产」部门（计件＋计时）＋手工分录',
             '制造': '车间费用（日记账）＋车间辅助工资＋车间折旧＋手工分录'}
_ITEM_NOTE = {'其中直接到款式': '填了「用在哪」款式的：等这个款交货才转出',
              '其中公共': '没写用在哪的：按交货双数×系数分',
              '公共上月结转': '上月没分完的（1 月＝年初在制，一般 0）',
              '公共本月可分摊': '上月结转来＋本月公共',
              '公共月末留在制': '下面「在制估计」按四块可分摊额拆到这块',
              '公共本月分摊': '有交货：可分摊−留在制；没交货：0（整月结转）',
              '公共结转下月': '可分摊−本月分到各款',
              '直接本月转出': '本月交货的款：累计直接成本−以前已转出',
              '本月转出合计': '直接转出＋公共分摊 → 月末结转营业成本'}


def build_alloc(wb, ctx):
    ws = wb[SH_ALLOC]
    last = AL_TOT
    widths(ws, {'A': 24, 'B': 40, **{al_col(m): 11.5 for m in range(1, 13)}, AL_TOT: 13, 'P': 10})
    title(ws, '成 本 分 摊 表（月度成本计算单 · 每个月各款式的成本是怎么来的）', last, C_VIEW,
          '💡 全自动，每个月一列（只有「在制估计」一行可以手填）。成本分四块：材料、外发加工、直接人工、制造费用，每块分两种：'
          '① 直接记到款式的（送货单 / 外发 / 计件工资 / 日记账里「用在哪」填了款式编码的）——这个款本月有交货才转出（到本月为止累计的减去以前转过的），'
          '没交货就挂着等交货；② 公共的（没写用在哪的材料、外发、计时工资、房租水电、车间折旧……）——本月有交货就按「交货双数 × 款式系数」'
          '全部分到本月交货的各款；一双都没交的月份整月结转到下个月。'
          '「在制估计」：月末车间里还有一大批没做完、不想让本月交货的款把公共成本全背了，就估个数填进去（比如 3000），这部分留到下个月再分；不填＝全部分掉。'
          '「本月转出合计」月末自动结转到营业成本；「月末在制」＝生产成本科目余额（挂着还没交货的）。最下面「核对」行 12 个月都要是 √。'
          '晚到的送货单录进去以后，那个月的数会变（【基础资料】结了账的月份除外）。', h2=80)
    home_link(ws, 'P1')
    R = AL_ROWS
    CR = AL_CROWS
    # 第 3 行：一眼看平不平
    _lbl(ws, 'A3', '全年核对')
    chk_all = f'{AL_TOT}{R["核对"]}'
    _val(ws, 'B3', f'=IF({chk_all}=0,"√ 12 个月都平：各款式成本加起来＝转出合计","✗ 有月份对不上：看最下面「核对」行")',
         None, 'B3:F3', F_KPI_L)
    _lbl(ws, 'G3', '本年转出\n（营业成本）', 'G3:H3')
    _val(ws, 'I3', f'={AL_TOT}{R["转出合计"]}', MONEY, 'I3:J3')
    _lbl(ws, 'K3', '12 月末在制\n（生产成本余额）', 'K3:L3')
    _val(ws, 'M3', f'={al_col(12)}{R["月末在制"]}', MONEY, 'M3:N3')
    ws.row_dimensions[3].height = 32
    ws.conditional_formatting.add('B3', FormulaRule(formula=['LEFT($B$3,1)="✗"'], fill=FILL_WARN, font=F_RED))
    # 表头
    header(ws, 4, [('A', '项目'), ('B', '说明')] + [(al_col(m), f'{m}月') for m in range(1, 13)] + [(AL_TOT, '全年')], H_VIEW, height=24)
    allm = [al_col(m) for m in range(1, 13)]
    C1, C12 = allm[0], allm[-1]

    def row_style(r, label, note, fmt=MONEY, bold=False, fill_=FILL_NONE):
        put(ws, f'A{r}', label, F_TXTB if bold else F_TXT, fill_, align=AL)
        put(ws, f'B{r}', note, F_NOTE, fill_, align=AL)
        for c in allm + [AL_TOT]:
            put(ws, f'{c}{r}', None, F_AUTOB if bold else F_AUTO, fill_, fmt, AR)

    def sum_year(r):
        ws[f'{AL_TOT}{r}'] = f'=ROUND(SUM({C1}{r}:{C12}{r}),2)'

    # 双数、加权双数
    dq_col = lambda m: smr('双数', m)
    row_style(R['双数'], '本月交货双数', '【订单明细】本月交货（只算实交>0、能记账的行）', INT, True, FILL_SUB)
    row_style(R['加权材料'], '加权双数（材料）', 'Σ 各款 交货双数×材料权重（分材料公共成本用）', WT)
    row_style(R['加权工费'], '加权双数（工费）', 'Σ 各款 交货双数×工费权重（分外发/人工/制造用）', WT)
    wm_r, wl_r = rng(SH_SM, SM_WM, SM_R0, SM_R1), rng(SH_SM, SM_WL, SM_R0, SM_R1)
    for m in range(1, 13):
        c = al_col(m)
        ws[f'{c}{R["双数"]}'] = f'=SUM({dq_col(m)})'
        ws[f'{c}{R["加权材料"]}'] = f'=SUMPRODUCT({dq_col(m)},{wm_r})'
        ws[f'{c}{R["加权工费"]}'] = f'=SUMPRODUCT({dq_col(m)},{wl_r})'
    for k in ('双数', '加权材料', '加权工费'):
        ws[f'{AL_TOT}{R[k]}'] = f'=SUM({C1}{R[k]}:{C12}{R[k]})'

    # 四块
    bands = {'材料': '① 材料', '外发': '② 外发加工', '人工': '③ 直接人工', '制造': '④ 制造费用（车间费用＋车间辅助工资＋车间折旧）'}
    band_color = {'材料': 'FF548235', '外发': 'FF2F75B5', '人工': 'FFBF8F00', '制造': 'FF7030A0'}
    for comp in COMPS:
        rr = CR[comp]
        section(ws, rr['本月发生'] - 1, 'A', AL_TOT, bands[comp], band_color[comp])
        for it in AL_ITEMS:
            r = rr[it]
            key_row = it in ('本月发生', '本月转出合计')
            note = _SRC_NOTE[comp] if it == '本月发生' else _ITEM_NOTE[it]
            if comp == '制造' and it == '其中公共':
                note = '没写用在哪的＋车间折旧：按交货双数×系数分'
            row_style(r, _ITEM_LBL[it].format(c=_COMP_NAME[comp]), note, MONEY, key_row,
                      FILL_SUB if it == '本月转出合计' else (fill('FFF7F7F7') if it == '本月发生' else FILL_NONE))
        wrow = _w(comp)[1]
        for m in range(1, 13):
            c = al_col(m)
            g = lambda it: f'{c}{rr[it]}'
            direct = cost_sum(f'"款|*|{comp}"', str(m), comp)
            pub = cost_sum(f'"公||{comp}"', str(m), comp) + (f'+{fa_dep(m, "车间")}' if comp == '制造' else '')
            ws[g('其中直接到款式')] = f'=ROUND({direct},2)'
            ws[g('其中公共')] = f'=ROUND({pub},2)'
            ws[g('本月发生')] = f'=ROUND({g("其中直接到款式")}+{g("其中公共")},2)'
            if m == 1:   # 1 月：年初在制（会计科目表 4001xx 年初余额，一般 0）
                ws[g('公共上月结转')] = f'=ROUND(SUMIF({COA_CODES},"{COMP_CODE[comp]}",{COA_OPENS}),2)'
            else:
                ws[g('公共上月结转')] = f'={al_col(m - 1)}{rr["公共结转下月"]}'
            ws[g('公共本月可分摊')] = f'=ROUND({g("公共上月结转")}+{g("其中公共")},2)'
            tot = '+'.join(f'{c}{CR[x]["公共本月可分摊"]}' for x in COMPS)
            est = f'N({c}{R["在制估计"]})'
            base = f'{c}{R[wrow]}'
            ws[g('公共月末留在制')] = (f'=IF(OR({est}<=0,({tot})<=0,{base}<=0),0,'
                                      f'ROUND({g("公共本月可分摊")}*MIN(1,{est}/({tot})),2))')
            ws[g('公共本月分摊')] = f'=IF({base}>0,ROUND({g("公共本月可分摊")}-{g("公共月末留在制")},2),0)'
            ws[g('公共结转下月')] = f'=ROUND({g("公共本月可分摊")}-{g("公共本月分摊")},2)'
            ws[g('直接本月转出')] = f'=ROUND(SUM({smr("转" + comp, m)}),2)'
            ws[g('本月转出合计')] = f'=ROUND({g("直接本月转出")}+{g("公共本月分摊")},2)'
        for it in ('本月发生', '其中直接到款式', '其中公共', '公共本月分摊', '直接本月转出', '本月转出合计'):
            sum_year(rr[it])
        ws[f'{AL_TOT}{rr["公共上月结转"]}'] = f'={C1}{rr["公共上月结转"]}'
        ws[f'{AL_TOT}{rr["公共本月可分摊"]}'] = f'=ROUND({AL_TOT}{rr["公共上月结转"]}+{AL_TOT}{rr["其中公共"]},2)'
        ws[f'{AL_TOT}{rr["公共月末留在制"]}'] = f'={C12}{rr["公共月末留在制"]}'
        ws[f'{AL_TOT}{rr["公共结转下月"]}'] = f'={C12}{rr["公共结转下月"]}'

    # 合计段
    section(ws, R['发生合计'] - 1, 'A', AL_TOT, '合计 · 在制', 'FF833C0C')
    row_style(R['发生合计'], '本月发生合计', '四块本月发生合计（＝本月记进生产成本的）', MONEY, True)
    row_style(R['转出合计'], '本月转出合计', '四块转出合计＝本月营业成本（自动结转）', MONEY, True, FILL_TOT)
    row_style(R['月末在制'], '月末在制（生产成本余额）', '年初在制＋累计发生−累计转出（挂着没交货的）', MONEY, True)
    row_style(R['在制估计'], '在制估计（手填，可空）', '月末车间没做完、要留到下月的公共成本估计数', MONEY)
    row_style(R['平均单双'], '平均每双成本', '本月转出合计÷本月交货双数', MONEY)
    section(ws, R['标准材料'] - 1, 'A', AL_TOT, '实际材料 vs 标准（只算【款式档案】填了「标准材料成本/双」的款）', 'FF7F7F7F')
    row_style(R['标准材料'], '标准材料成本', 'Σ 交货双数×标准材料/双（填了标准的款）', MONEY)
    row_style(R['实际材料'], '实际材料成本', '这些款本月分到的材料成本（直接＋分摊）', MONEY)
    row_style(R['材料差额'], '差额（实际−标准）', '负数＝比标准少', MONEY)
    row_style(R['提示'], '提示', '比标准少 10% 以上：可能还有送货单没拿到', 'General')
    row_style(R['核对'], '核对', 'Σ 各款式本月成本 − 本月转出合计，应为 0（√＝平）', CHKF, True)
    sdm = str_(ST_SDM)
    occ_rows = [CR[x]['本月发生'] for x in COMPS]
    out_rows = [CR[x]['本月转出合计'] for x in COMPS]
    for m in range(1, 13):
        c = al_col(m)
        ws[f'{c}{R["发生合计"]}'] = '=ROUND(' + '+'.join(f'{c}{r}' for r in occ_rows) + ',2)'
        ws[f'{c}{R["转出合计"]}'] = '=ROUND(' + '+'.join(f'{c}{r}' for r in out_rows) + ',2)'
        open_wip = '+'.join(f'${C1}${CR[x]["公共上月结转"]}' for x in COMPS)
        ws[f'{c}{R["月末在制"]}'] = (f'=ROUND({open_wip}+SUM(${C1}{R["发生合计"]}:{c}{R["发生合计"]})'
                                    f'-SUM(${C1}{R["转出合计"]}:{c}{R["转出合计"]}),2)')
        ws[f'{c}{R["平均单双"]}'] = f'=IF({c}{R["双数"]}>0,ROUND({c}{R["转出合计"]}/{c}{R["双数"]},2),0)'
        ws[f'{c}{R["标准材料"]}'] = f'=ROUND(SUMPRODUCT({dq_col(m)},{sdm}),2)'
        ws[f'{c}{R["实际材料"]}'] = f'=ROUND(SUMPRODUCT({smr("本材料", m)},ISNUMBER({sdm})*({sdm}>0)),2)'
        ws[f'{c}{R["材料差额"]}'] = f'=IF({c}{R["标准材料"]}=0,0,ROUND({c}{R["实际材料"]}-{c}{R["标准材料"]},2))'
        ws[f'{c}{R["核对"]}'] = f'=ROUND(SUM({smr("成本", m)})-{c}{R["转出合计"]},2)'
        ws[f'{c}{R["在制估计"]}'].fill = FILL_IN
        ws[f'{c}{R["在制估计"]}'].font = F_IN
    for c in allm + [AL_TOT]:
        s, d = f'{c}{R["标准材料"]}', f'{c}{R["材料差额"]}'
        ws[f'{c}{R["提示"]}'] = f'=IF(AND({s}>0,{d}<-0.1*{s}),"⚠ 比标准少 "&TEXT(-{d}/{s},"0%"),"")'
        ws[f'{c}{R["提示"]}'].alignment = AC
    for k in ('发生合计', '转出合计', '标准材料', '实际材料'):
        sum_year(R[k])
    T = AL_TOT
    ws[f'{T}{R["月末在制"]}'] = f'={C12}{R["月末在制"]}'
    ws[f'{T}{R["平均单双"]}'] = f'=IF({T}{R["双数"]}>0,ROUND({T}{R["转出合计"]}/{T}{R["双数"]},2),0)'
    ws[f'{T}{R["材料差额"]}'] = f'=IF({T}{R["标准材料"]}=0,0,ROUND({T}{R["实际材料"]}-{T}{R["标准材料"]},2))'
    ws[f'{T}{R["核对"]}'] = f'=ROUND(SUMPRODUCT(ABS({C1}{R["核对"]}:{C12}{R["核对"]})),2)'
    ws[f'{T}{R["在制估计"]}'].fill = FILL_AUTO
    # 在制估计：只能填 ≥0 的数
    dv = DataValidation(type='decimal', operator='greaterThanOrEqual', formula1='0', allow_blank=True, showErrorMessage=True,
                        errorTitle='在制估计', error='填一个金额（≥0），不想留就空着')
    dv.promptTitle, dv.prompt, dv.showInputMessage = '在制估计', '月末车间没做完的、要留到下个月再分的公共成本，估个数；不填＝全部分掉', True
    ws.add_data_validation(dv)
    dv.add(f'{C1}{R["在制估计"]}:{C12}{R["在制估计"]}')
    # 条件格式：核对不为 0 红；提示黄；负数红字
    rk = R['核对']
    ws.conditional_formatting.add(f'{C1}{rk}:{T}{rk}', FormulaRule(formula=[f'ROUND({C1}{rk},2)<>0'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(f'{C1}{rk}:{T}{rk}', FormulaRule(formula=[f'ROUND({C1}{rk},2)=0'], font=F_GOOD))
    rt = R['提示']
    ws.conditional_formatting.add(f'{C1}{rt}:{T}{rt}', FormulaRule(formula=[f'LEFT({C1}{rt},1)="⚠"'], fill=FILL_YEL))
    for r in range(4, rk + 1):
        ws.row_dimensions[r].height = 18
    ws.row_dimensions[4].height = 24
    ws.freeze_panes = 'C5'
    print_setup(ws, '4:4', landscape=True)


# ═══════════════════════════ 款式成本利润 ═══════════════════════════
SPL_R0 = 6                      # 清单第一行（第 4 行表头、第 5 行合计）
SPL_N = ST_N                    # 最多 400 个款
SPL_R1 = SPL_R0 + SPL_N - 1
# 隐藏帮手列：AF3 单款行号；AG3 起始月 AH3 截止月 AI3 期间费用 AJ3 全部收入；
#   AG～AQ（第 6～405 行，第 k 行＝第 k 个款）：编码、双数、收入、材料、外发、人工、制造、成本、订单数、序号键、第 k 个是第几个款
#   AS、AT（第 5～6004 行，跟【订单明细】一行对一行）：「订单|款式」在所选月份第一次交货＝1；订单在所选月份第一次交货＝1
SPL_SROW = '$AF$3'
SPL_LO, SPL_HI, SPL_PEXP, SPL_REVT = '$AG$3', '$AH$3', '$AI$3', '$AJ$3'
(SH_CODE, SH_DQ, SH_REV, SH_MAT_, SH_OUT_, SH_LAB, SH_MOH, SH_COST, SH_NORD, SH_KEY, SH_IDX) = \
    'AG AH AI AJ AK AL AM AN AO AP AQ'.split()
SH_OSF, SH_OF = 'AS', 'AT'


def _je_net(code, lo, hi):
    """记账分录里某科目（前缀）在 lo～hi 月的「借−贷」发生额"""
    def s(col):
        return (f'SUMIFS({jer(JE_AMT)},{jer(col)},"{code}*",{jer(JE_M)},">="&{lo},{jer(JE_M)},"<="&{hi})')
    return f'({s(JE_DR)}-{s(JE_CR)})'


def build_spl(wb, ctx):
    ws = wb[SH_SPL]
    widths(ws, {'A': 5, 'B': 11, 'C': 12, 'D': 8, 'E': 12, 'F': 9, 'G': 11, 'H': 10, 'I': 10, 'J': 10, 'K': 12, 'L': 9,
                'M': 9, 'N': 12, 'O': 8, 'P': 11, 'Q': 12, 'R': 7, 'S': 2,
                'T': 7, 'U': 8, 'V': 11, 'W': 10, 'X': 10, 'Y': 10, 'Z': 10, 'AA': 11, 'AB': 9, 'AC': 11, 'AD': 8, 'AE': 2})
    title(ws, '款 式 成 本 利 润（每个款卖了多少 · 成本多少 · 赚多少）', 'AD', C_VIEW,
          '💡 老板看。黄格子选月份范围（起始月～截止月，1～12）。成本＝直接记到这个款的（专用材料、外发、计件工资、专用模具）'
          '＋按「交货双数×系数」分到的公共成本（没写用在哪的材料、计时工资、房租水电、车间折旧……），每个月怎么分的看【成本分摊表】。'
          '只列所选月份有交货（或有成本）的款；「只看有交货」选「是」就只列有交货双数的。'
          '「期间费用分摊」＝所选月份的管理、销售、财务费用按收入比例摊给各款（【基础资料】「期间费用摊到款式」可改成「不摊」），毛利减掉它就是「净利」。'
          '晚到的送货单录进去以后，那个月的成本会变（【基础资料】结了账的月份不变）。右边「单款逐月」选一个款看 1～12 月。', h2=64)
    home_link(ws, 'AE1')
    ws.column_dimensions['AE'].width = 10
    # 第 3 行：选择
    selector(ws, 'B3', '起始月', 'C3', 1, '"1,2,3,4,5,6,7,8,9,10,11,12"', MONTHF, '从几月（1～12）')
    selector(ws, 'D3', '截止月', 'E3', 12, '"1,2,3,4,5,6,7,8,9,10,11,12"', MONTHF, '到几月（1～12）')
    selector(ws, 'F3', '只看有交货', 'G3', '否', '"否,是"', None, '是＝只列所选月份有交货双数的款')
    _lbl(ws, 'H3', '所选月份\n期间费用', 'H3:I3')
    _val(ws, 'J3', f'={SPL_PEXP}', MONEY, 'J3:K3')
    _lbl(ws, 'L3', '跟成本分摊表', 'L3:M3')
    mrow = f'{SMQ}!${sm_col("双数", 1)}${SM_MROW}:${sm_col("双数", 12)}${SM_MROW}'
    out_row = f'{q(SH_ALLOC)}!${al_col(1)}${AL_ROWS["转出合计"]}:${al_col(12)}${AL_ROWS["转出合计"]}'
    alloc_out = f'SUMPRODUCT(({mrow}>={SPL_LO})*({mrow}<={SPL_HI})*{out_row})'
    _val(ws, 'N3', f'=IF(ROUND(SUM({SH_COST}{SPL_R0}:{SH_COST}{SPL_R1})-{alloc_out},2)=0,"√ 成本对得上",'
                   f'"✗ 差 "&TEXT(SUM({SH_COST}{SPL_R0}:{SH_COST}{SPL_R1})-{alloc_out},"#,##0.00"))', None, 'N3:O3', F_KPI_L)
    put(ws, 'P3', f'=IF({PERIOD}="按收入","期间费用按收入比例摊到各款","期间费用不摊（只算到毛利）")', F_NOTE, align=ALW, border=False)
    ws.merge_cells('P3:R3')
    ws.row_dimensions[3].height = 32
    ws.conditional_formatting.add('N3', FormulaRule(formula=['LEFT($N$3,1)="✗"'], fill=FILL_WARN, font=F_RED))
    # 隐藏的选择换算
    a = f'IF(ISNUMBER($C$3),MAX(1,MIN(12,INT($C$3))),1)'
    b = f'IF(ISNUMBER($E$3),MAX(1,MIN(12,INT($E$3))),12)'
    ws[SPL_LO.replace('$', '')] = f'=MIN({a},{b})'
    ws[SPL_HI.replace('$', '')] = f'=MAX({a},{b})'
    pexp = '+'.join(_je_net(code, SPL_LO, SPL_HI) for code in ('5601', '5602', '5603'))
    ws[SPL_PEXP.replace('$', '')] = f'=ROUND({pexp},2)'
    ws[SPL_REVT.replace('$', '')] = f'=ROUND(SUM({SH_REV}{SPL_R0}:{SH_REV}{SPL_R1}),2)'
    ws['AF2'], ws['AG2'], ws['AH2'], ws['AI2'], ws['AJ2'] = '单款行号', '起始月', '截止月', '期间费用', '全部收入'
    # 表头
    heads = [('A', '序号'), ('B', '款式编码'), ('C', '名称'), ('D', '交货\n双数'), ('E', '交货收入'), ('F', '平均\n单价'),
             ('G', '材料'), ('H', '外发加工'), ('I', '直接人工'), ('J', '制造费用'), ('K', '成本合计'), ('L', '单双\n成本'),
             ('M', '标准\n单双成本'), ('N', '毛利'), ('O', '毛利率'), ('P', '期间费用\n分摊'), ('Q', '净利'), ('R', '订单数')]
    header(ws, 4, heads, H_VIEW)
    for c, _ in heads[6:12]:
        ws[f'{c}4'].fill = fill('FF2F75B5')
    # 每个款在所选月份的数（帮手列，第 k 行＝第 k 个款）
    sm_dm = lambda blk, sr: (f'SUMPRODUCT(({SMQ}!${sm_col(blk, 1)}${SM_MROW}:${sm_col(blk, 12)}${SM_MROW}>={SPL_LO})'
                             f'*({SMQ}!${sm_col(blk, 1)}${SM_MROW}:${sm_col(blk, 12)}${SM_MROW}<={SPL_HI})'
                             f'*{SMQ}!${sm_col(blk, 1)}{sr}:${sm_col(blk, 12)}{sr})')
    only = '$G$3="是"'
    for r in range(SPL_R0, SPL_R1 + 1):
        sr = SM_R0 + (r - SPL_R0)                   # _款式月 / 款式档案 的行
        e = f'${SH_CODE}{r}=""'
        ws[f'{SH_CODE}{r}'] = f'={SMQ}!${SM_CODE}{sr}&""'
        for col, blk in ((SH_DQ, '双数'), (SH_REV, '收入'), (SH_MAT_, '本材料'), (SH_OUT_, '本外发'), (SH_LAB, '本人工'), (SH_MOH, '本制造')):
            ws[f'{col}{r}'] = f'=IF({e},0,ROUND({sm_dm(blk, sr)},2))'
        ws[f'{SH_COST}{r}'] = f'=ROUND({SH_MAT_}{r}+{SH_OUT_}{r}+{SH_LAB}{r}+{SH_MOH}{r},2)'
        ws[f'{SH_NORD}{r}'] = f'=IF({e},0,SUMIFS(${SH_OSF}${OD_R0}:${SH_OSF}${OD_R1},{odr(OD_SROW)},ROW()-{SPL_R0 - 1}))'
        ws[f'{SH_KEY}{r}'] = (f'=IF({e},"",IF(IF({only},{SH_DQ}{r}>0,OR({SH_DQ}{r}<>0,{SH_REV}{r}<>0,{SH_COST}{r}<>0)),'
                              f'ROW()-{SPL_R0 - 1},""))')
        ws[f'{SH_IDX}{r}'] = f'=IFERROR(SMALL(${SH_KEY}${SPL_R0}:${SH_KEY}${SPL_R1},ROW()-{SPL_R0 - 1}),0)'
    # 跟【订单明细】一行对一行：所选月份里「订单|款式」、订单第一次交货（数订单个数用）
    for r in range(OD_R0, OD_R1 + 1):
        ok, dq, dm, osk = _od(OD_OK, r), _od(OD_DQ, r), _od(OD_DM, r), _od(OD_OSKEY, r)
        hit = f'AND({ok}=1,{dq}>0,{dm}>={SPL_LO},{dm}<={SPL_HI})'
        rest = (f'{_odg(OD_OK, r - 1)},1,{_odg(OD_DQ, r - 1)},">0",{_odg(OD_DM, r - 1)},">="&{SPL_LO},'
                f'{_odg(OD_DM, r - 1)},"<="&{SPL_HI}')
        ws[f'{SH_OSF}{r}'] = f'=IF({hit},IF(COUNTIFS({_odg(OD_OSKEY, r - 1)},{esc(osk)},{rest})=0,1,0),0)'
        left = 'LEFT(' + osk + ',FIND("|",' + osk + '))'
        pre = esc(left) + '&"*"'
        ws[f'{SH_OF}{r}'] = f'=IF({hit},IF(COUNTIFS({_odg(OD_OSKEY, r - 1)},{pre},{rest})=0,1,0),0)'
    # 清单
    rg = lambda col: f'${col}${SPL_R0}:${col}${SPL_R1}'
    for r in range(SPL_R0, SPL_R1 + 1):
        ix = f'${SH_IDX}{r}'
        g = lambda col: f'INDEX({rg(col)},{ix})'
        e = f'{ix}=0'
        ws[f'A{r}'] = f'=IF({e},"",ROW()-{SPL_R0 - 1})'
        ws[f'B{r}'] = f'=IF({e},"",{g(SH_CODE)})'
        ws[f'C{r}'] = f'=IF({e},"",INDEX({str_(ST_NAME)},{ix})&"")'
        for col, src in (('D', SH_DQ), ('E', SH_REV), ('G', SH_MAT_), ('H', SH_OUT_), ('I', SH_LAB), ('J', SH_MOH)):
            ws[f'{col}{r}'] = f'=IF({e},"",{g(src)})'
        ws[f'F{r}'] = f'=IF({e},"",IF(N(D{r})>0,ROUND(E{r}/D{r},2),""))'
        ws[f'K{r}'] = f'=IF({e},"",ROUND(G{r}+H{r}+I{r}+J{r},2))'
        ws[f'L{r}'] = f'=IF({e},"",IF(N(D{r})>0,ROUND(K{r}/D{r},2),""))'
        sdm, sdl = f'INDEX({str_(ST_SDM)},{ix})', f'INDEX({str_(ST_SDL)},{ix})'
        ws[f'M{r}'] = f'=IF({e},"",IF(AND(N({sdm})=0,N({sdl})=0),"",N({sdm})+N({sdl})))'
        ws[f'N{r}'] = f'=IF({e},"",ROUND(E{r}-K{r},2))'
        ws[f'O{r}'] = f'=IF({e},"",IF(N(E{r})<=0,"",ROUND(N{r}/E{r},4)))'
        ws[f'P{r}'] = f'=IF({e},"",IF({PERIOD}<>"按收入","",IF({SPL_REVT}=0,0,ROUND({SPL_PEXP}*E{r}/{SPL_REVT},2))))'
        ws[f'Q{r}'] = f'=IF({e},"",IF({PERIOD}<>"按收入","",ROUND(N{r}-P{r},2)))'
        ws[f'R{r}'] = f'=IF({e},"",{g(SH_NORD)})'
    cols = [CL(i) for i in range(1, 19)]
    style_rows(ws, SPL_R0, SPL_R1, cols, auto=cols,
               fmts={'D': INT, 'E': MONEY, 'F': MONEY, 'G': MONEY, 'H': MONEY, 'I': MONEY, 'J': MONEY, 'K': MONEY, 'L': MONEY,
                     'M': MONEY, 'N': MONEY, 'O': PCT, 'P': MONEY, 'Q': MONEY, 'R': INT},
               aligns={'B': AL, 'C': AL, **{c: AR for c in 'EFGHIJKLMNPQ'}}, bold=['K', 'N', 'Q'])
    for r in range(SPL_R0, SPL_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    # 合计行（第 5 行）
    put(ws, 'A5', '', F_TXTB, FILL_TOT)
    put(ws, 'B5', f'="合计（"&COUNT($A${SPL_R0}:$A${SPL_R1})&" 个款）"', F_TXTB, FILL_TOT, align=AC)
    ws.merge_cells('B5:C5')
    for c in 'DEGHIJKNPQ':
        put(ws, f'{c}5', f'=ROUND(SUM({c}{SPL_R0}:{c}{SPL_R1}),2)', F_AUTOB, FILL_TOT, INT if c == 'D' else MONEY, AR)
    put(ws, 'F5', f'=IF(D5>0,ROUND(E5/D5,2),"")', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'L5', f'=IF(D5>0,ROUND(K5/D5,2),"")', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'M5', '', F_AUTOB, FILL_TOT)
    put(ws, 'O5', f'=IF(E5<=0,"",ROUND(N5/E5,4))', F_AUTOB, FILL_TOT, PCT, AR)
    put(ws, 'R5', f'=SUM(${SH_OF}${OD_R0}:${SH_OF}${OD_R1})', F_AUTOB, FILL_TOT, INT, AC)
    ws['P5'] = f'=IF({PERIOD}<>"按收入","",ROUND(SUM(P{SPL_R0}:P{SPL_R1}),2))'
    ws['Q5'] = f'=IF({PERIOD}<>"按收入","",ROUND(SUM(Q{SPL_R0}:Q{SPL_R1}),2))'
    ws.row_dimensions[5].height = 22
    for c in 'NQ':
        _neg_red(ws, f'{c}5:{c}{SPL_R1}', f'{c}5')

    # 右边：单款逐月
    put(ws, 'T3', '单款逐月', F_KPI_L, fill('FFD9E1F2'), align=AC)
    first = next((x[ST_CODE] for x in ctx.get(SH_STY, []) if x.get(ST_CODE) == 'HK09'), None) or \
        next((x[ST_CODE] for x in ctx.get(SH_STY, []) if x.get(ST_CODE)), None)
    put(ws, 'U3', first, F_SEL, FILL_SEL, align=AC)
    ws.merge_cells('U3:V3')
    dv = DataValidation(type='list', formula1=f'={ST_CODES}', allow_blank=True, showErrorMessage=False)
    dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', '选一个款式，看它 1～12 月', True
    ws.add_data_validation(dv)
    dv.add('U3')
    ws[SPL_SROW.replace('$', '')] = f'=IF(TRIM($U$3&"")="",0,IFERROR(MATCH(TRIM($U$3&""),{ST_KEYS},0),0))'
    put(ws, 'W3', f'=IF({SPL_SROW}=0,IF(TRIM($U$3&"")="","← 选款式","✗ 这个款式没在【款式档案】登记"),'
                  f'INDEX({str_(ST_NAME)},{SPL_SROW})&"　结算单价 "&TEXT(N(INDEX({str_(ST_PRICE)},{SPL_SROW})),"0.00")&" 元/双")',
        F_NOTE, align=AL, border=False)
    ws.merge_cells('W3:AD3')
    sh = [('T', '月份'), ('U', '交货\n双数'), ('V', '交货收入'), ('W', '材料'), ('X', '外发加工'), ('Y', '直接人工'), ('Z', '制造费用'),
          ('AA', '成本合计'), ('AB', '单双\n成本'), ('AC', '毛利'), ('AD', '毛利率')]
    header(ws, 4, sh, 'FF2F75B5')
    blocks = {'U': '双数', 'V': '收入', 'W': '本材料', 'X': '本外发', 'Y': '本人工', 'Z': '本制造', 'AA': '成本'}
    for m in range(1, 13):
        r = SPL_R0 + m - 1
        put(ws, f'T{r}', m, F_TXTB, FILL_SUB, MONTHF, AC)
        for col, blk in blocks.items():
            two_d = f'{SMQ}!${sm_col(blk, 1)}${SM_R0}:${sm_col(blk, 12)}${SM_R1}'
            put(ws, f'{col}{r}', f'=IF({SPL_SROW}=0,"",INDEX({two_d},{SPL_SROW},$T{r}))', F_AUTO, None,
                INT if col == 'U' else MONEY, AR)
        put(ws, f'AB{r}', f'=IF({SPL_SROW}=0,"",IF(N(U{r})>0,ROUND(AA{r}/U{r},2),""))', F_AUTO, None, MONEY, AR)
        put(ws, f'AC{r}', f'=IF({SPL_SROW}=0,"",ROUND(V{r}-AA{r},2))', F_AUTOB, None, MONEY, AR)
        put(ws, f'AD{r}', f'=IF({SPL_SROW}=0,"",IF(N(V{r})<=0,"",ROUND(AC{r}/V{r},4)))', F_AUTO, None, PCT, AR)
    rN = SPL_R0 + 11
    put(ws, 'T5', '全年', F_TXTB, FILL_TOT, align=AC)
    for col in list(blocks) + ['AC']:
        put(ws, f'{col}5', f'=IF({SPL_SROW}=0,"",ROUND(SUM({col}{SPL_R0}:{col}{rN}),2))', F_AUTOB, FILL_TOT,
            INT if col == 'U' else MONEY, AR)
    put(ws, 'AB5', f'=IF({SPL_SROW}=0,"",IF(N(U5)>0,ROUND(AA5/U5,2),""))', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'AD5', f'=IF({SPL_SROW}=0,"",IF(N(V5)<=0,"",ROUND(AC5/V5,4)))', F_AUTOB, FILL_TOT, PCT, AR)
    _neg_red(ws, f'AC5:AC{rN}', 'AC5')
    # 帮手列隐藏
    _help_cols(ws, 4, [(SH_CODE, '编码'), (SH_DQ, '双数'), (SH_REV, '收入'), (SH_MAT_, '材料'), (SH_OUT_, '外发'),
                       (SH_LAB, '人工'), (SH_MOH, '制造'), (SH_COST, '成本'), (SH_NORD, '订单数'), (SH_KEY, '序号键'),
                       (SH_IDX, '第k个'), (SH_OSF, '订单|款式首次'), (SH_OF, '订单首次')], SPL_R0, SPL_R1)
    for r in range(OD_R0, OD_R1 + 1):
        ws[f'{SH_OSF}{r}'].font = ws[f'{SH_OF}{r}'].font = F_HELP
    for c in ('AF', 'AG', 'AH', 'AI', 'AJ'):
        ws[f'{c}2'].font = ws[f'{c}3'].font = F_HELP
    hide(ws, 'AF', 'AR')
    ws.freeze_panes = f'C{SPL_R0}'
    print_setup(ws, '4:5', landscape=True)
    ws.print_area = f'A1:AD{SPL_R1}'


# ═══════════════════════════ 订单汇总 ═══════════════════════════
OS_R0, OS_N = 6, 600
OS_R1 = OS_R0 + OS_N - 1
OS_IDX, OS_DMIN, OS_DMAX, OS_OMIN = 'T', 'U', 'V', 'W'       # 隐藏：第 k 张单在【订单明细】第几行；首次/最后交货、下单日期的原始数
OS_HD, OS_HO = 'Y', 'Z'                                      # 隐藏（跟【订单明细】一行对一行）：交货日期数字、下单日期数字


def build_osum(wb, ctx):
    ws = wb[SH_OSUM]
    widths(ws, {'A': 5, 'B': 11, 'C': 10, 'D': 11, 'E': 7, 'F': 7, 'G': 8, 'H': 8, 'I': 8, 'J': 8, 'K': 11, 'L': 11,
                'M': 12, 'N': 12, 'O': 12, 'P': 8, 'Q': 8, 'R': 30, 'S': 2})
    title(ws, '订 单 汇 总（每张订单 · 交货进度 · 收入成本毛利）', 'R', C_VIEW,
          '💡 每张订单（电商部采购单）一行，按【订单明细】里第一次出现的顺序列出。交货进度＝已交双数÷订单双数；'
          '订单的成本＝每行 实交双数×那个款那个月的单双成本（【成本分摊表】分出来的）；没交的行没有成本；'
          '款式没在【款式档案】登记的行也没有成本（「提示」列会说）。退回来的（实交填负数）冲减收入，不算在已交双数里。'
          '想看某一张单的每一行，去【订单查询】选订单号。', h2=48)
    home_link(ws, 'S1')
    ws.column_dimensions['S'].width = 10
    rgv = lambda col: f'${col}${OS_R0}:${col}${OS_R1}'
    _lbl(ws, 'A3', '订单数', 'A3:B3')
    _val(ws, 'C3', f'=COUNT({odr(OD_OFIRST)})', INT)
    _lbl(ws, 'D3', '已交完')
    _val(ws, 'E3', f'=COUNTIF({rgv("Q")},"已交完")', INT, 'E3:F3')
    _lbl(ws, 'G3', '部分交')
    _val(ws, 'H3', f'=COUNTIF({rgv("Q")},"部分交")', INT)
    _lbl(ws, 'I3', '未交')
    _val(ws, 'J3', f'=COUNTIF({rgv("Q")},"未交")', INT)
    put(ws, 'K3', f'=IF(C3>{OS_N},"⚠ 订单超过 {OS_N} 张，只列前 {OS_N} 张","")', F_RED, align=AL, border=False)
    ws.merge_cells('K3:R3')
    ws.row_dimensions[3].height = 30
    heads = [('A', '序号'), ('B', '订单号'), ('C', '客户'), ('D', '下单日期'), ('E', '款式数'), ('F', '订单\n行数'),
             ('G', '订单\n双数'), ('H', '已交\n双数'), ('I', '未交\n双数'), ('J', '交货\n进度'), ('K', '首次交货'), ('L', '最后交货'),
             ('M', '交货收入'), ('N', '成本'), ('O', '毛利'), ('P', '毛利率'), ('Q', '状态'), ('R', '提示')]
    header(ws, 4, heads, H_VIEW)
    # 跟【订单明细】一行对一行：交货日期、下单日期变成数字（取最早/最晚用，文字、空的当 0）
    for r in range(OD_R0, OD_R1 + 1):
        ws[f'{OS_HD}{r}'] = f'=IF({_od(OD_DQ, r)}<>0,INT({_od(OD_DDATE, r)}),0)'
        ws[f'{OS_HO}{r}'] = f'=IF(ISNUMBER({_od(OD_DATE, r)}),INT({_od(OD_DATE, r)}),0)'
    hd, ho = f'${OS_HD}${OD_R0}:${OS_HD}${OD_R1}', f'${OS_HO}${OD_R0}:${OS_HO}${OD_R1}'
    no = odr(OD_NO)
    for r in range(OS_R0, OS_R1 + 1):
        ix = f'${OS_IDX}{r}'
        e = f'{ix}=0'
        B = f'$B{r}'
        crit = no + ',' + esc(B + '&""')
        ws[f'{OS_IDX}{r}'] = f'=IFERROR(SMALL({odr(OD_OFIRST)},ROW()-{OS_R0 - 1}),0)'
        ws[f'{OS_DMAX}{r}'] = f'=IF({e},0,SUMPRODUCT(MAX(({no}={B})*{hd})))'
        ws[f'{OS_DMIN}{r}'] = f'=IF({e},0,SUMPRODUCT(MAX(({no}={B})*({hd}>0)*(100000000-{hd}))))'
        ws[f'{OS_OMIN}{r}'] = f'=IF({e},0,SUMPRODUCT(MAX(({no}={B})*({ho}>0)*(100000000-{ho}))))'
        ws[f'A{r}'] = f'=IF({e},"",ROW()-{OS_R0 - 1})'
        ws[f'B{r}'] = f'=IF({e},"",INDEX({no},{ix}))'
        ws[f'C{r}'] = f'=IF({e},"",INDEX({odr(OD_CUSTK)},{ix})&"")'
        ws[f'D{r}'] = f'=IF({e},"",IF(${OS_OMIN}{r}=0,"",100000000-${OS_OMIN}{r}))'
        ws[f'E{r}'] = f'=IF({e},"",SUMIFS({odr(OD_OSFIRST)},{crit}))'
        ws[f'F{r}'] = f'=IF({e},"",COUNTIFS({crit}))'
        ws[f'G{r}'] = f'=IF({e},"",SUMIFS({odr(OD_OQ)},{crit}))'
        ws[f'H{r}'] = f'=IF({e},"",SUMIFS({odr(OD_DQ)},{crit},{odr(OD_DQ)},">0"))'
        ws[f'I{r}'] = f'=IF({e},"",MAX(0,G{r}-H{r}))'
        ws[f'J{r}'] = f'=IF({e},"",IF(N(G{r})>0,ROUND(H{r}/G{r},4),""))'
        ws[f'K{r}'] = f'=IF({e},"",IF(${OS_DMIN}{r}=0,"",100000000-${OS_DMIN}{r}))'
        ws[f'L{r}'] = f'=IF({e},"",IF(${OS_DMAX}{r}=0,"",${OS_DMAX}{r}))'
        ws[f'M{r}'] = f'=IF({e},"",ROUND(SUMIFS({odr(OD_AMT)},{crit},{odr(OD_OK)},1),2))'
        ws[f'N{r}'] = f'=IF({e},"",ROUND(SUMIFS({odr(OD_COSTV)},{crit}),2))'
        ws[f'O{r}'] = f'=IF({e},"",ROUND(M{r}-N{r},2))'
        ws[f'P{r}'] = f'=IF({e},"",IF(N(M{r})<=0,"",ROUND(O{r}/M{r},4)))'
        ws[f'Q{r}'] = f'=IF({e},"",IF(N(H{r})<=0,"未交",IF(H{r}<G{r},"部分交","已交完")))'
        nx = f'COUNTIFS({crit},{odr(OD_CHK)},"✗*")'
        nu = f'COUNTIFS({crit},{odr(OD_SROW)},0)'
        ws[f'R{r}'] = (f'=IF({e},"",IF({nx}>0,"✗ "&{nx}&" 行要改（看【订单明细】校验）",'
                       f'IF({nu}>0,"⚠ "&{nu}&" 行款式没登记：没有成本","")))')
    cols = [CL(i) for i in range(1, 19)]
    style_rows(ws, OS_R0, OS_R1, cols, auto=cols,
               fmts={'D': DATE, 'E': INT, 'F': INT, 'G': INT, 'H': INT, 'I': INT, 'J': PCT, 'K': DATE, 'L': DATE,
                     'M': MONEY, 'N': MONEY, 'O': MONEY, 'P': PCT},
               aligns={'B': AL, 'C': AL, 'R': AL, 'M': AR, 'N': AR, 'O': AR}, bold=['B', 'O'])
    for r in range(OS_R0, OS_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    # 合计行
    put(ws, 'A5', '', F_TXTB, FILL_TOT)
    put(ws, 'B5', f'="合计（"&COUNT({rgv("A")})&" 张）"', F_TXTB, FILL_TOT, align=AC)
    ws.merge_cells('B5:D5')
    for c in 'EFGHI':
        put(ws, f'{c}5', f'=SUM({c}{OS_R0}:{c}{OS_R1})', F_AUTOB, FILL_TOT, INT, AC)
    put(ws, 'J5', '=IF(G5>0,ROUND(H5/G5,4),"")', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, 'K5', f'=IF(COUNT({rgv("K")})=0,"",MIN({rgv("K")}))', F_AUTOB, FILL_TOT, DATE, AC)
    put(ws, 'L5', f'=IF(COUNT({rgv("L")})=0,"",MAX({rgv("L")}))', F_AUTOB, FILL_TOT, DATE, AC)
    for c in 'MNO':
        put(ws, f'{c}5', f'=ROUND(SUM({c}{OS_R0}:{c}{OS_R1}),2)', F_AUTOB, FILL_TOT, MONEY, AR)
    put(ws, 'P5', '=IF(M5<=0,"",ROUND(O5/M5,4))', F_AUTOB, FILL_TOT, PCT, AC)
    put(ws, 'Q5', f'="未交完 "&(COUNTIF({rgv("Q")},"未交")+COUNTIF({rgv("Q")},"部分交"))&" 张"', F_AUTOB, FILL_TOT, align=AC)
    put(ws, 'R5', f'=IF(COUNTIF({rgv("R")},"✗*")+COUNTIF({rgv("R")},"⚠*")=0,"",'
                  f'"有提示的 "&(COUNTIF({rgv("R")},"✗*")+COUNTIF({rgv("R")},"⚠*"))&" 张")', F_AUTOB, FILL_TOT, align=AL)
    ws.row_dimensions[5].height = 22
    _neg_red(ws, f'O5:O{OS_R1}', 'O5')
    _chk_cf(ws, f'R{OS_R0}:R{OS_R1}', f'$R{OS_R0}')
    ws.conditional_formatting.add(f'Q{OS_R0}:Q{OS_R1}', FormulaRule(formula=[f'$Q{OS_R0}="已交完"'], font=F_GOOD))
    ws.conditional_formatting.add(f'Q{OS_R0}:Q{OS_R1}', FormulaRule(formula=[f'$Q{OS_R0}="部分交"'], fill=FILL_YEL))
    ws.conditional_formatting.add(f'Q{OS_R0}:Q{OS_R1}', FormulaRule(formula=[f'$Q{OS_R0}="未交"'], font=F_WAIT))
    _help_cols(ws, 4, [(OS_IDX, '订单明细第几行'), (OS_DMIN, '首次交货(反)'), (OS_DMAX, '最后交货'), (OS_OMIN, '下单日期(反)')],
               OS_R0, OS_R1)
    _help_cols(ws, 4, [(OS_HD, '交货日期数字'), (OS_HO, '下单日期数字')], OD_R0, OD_R1)
    hide(ws, 'X')
    ws.auto_filter.ref = f'A4:R{OS_R1}'
    ws.freeze_panes = f'C{OS_R0}'
    print_setup(ws, '4:5', landscape=True)
    ws.print_area = f'A1:R{OS_R1}'


# ═══════════════════════════ 订单查询 ═══════════════════════════
OQ_R0, OQ_N = 8, 200
OQ_R1 = OQ_R0 + OQ_N - 1
OQ_SN = 100                                   # 按款式小计最多几个款
OQ_IDX, OQ_ODATE, OQ_SF, OQ_SIDX = 'Y', 'Z', 'AA', 'AB'


def build_oq(wb, ctx):
    ws = wb[SH_OQ]
    widths(ws, {'A': 5, 'B': 11, 'C': 14, 'D': 9, 'E': 8, 'F': 11, 'G': 7, 'H': 9, 'I': 11, 'J': 9, 'K': 11, 'L': 11,
                'M': 8, 'N': 18, 'O': 30, 'P': 2, 'Q': 11, 'R': 8, 'S': 8, 'T': 11, 'U': 11, 'V': 11, 'W': 8, 'X': 2})
    title(ws, '订 单 查 询（选一张订单 · 看每一行交了多少、收入成本毛利）', 'W', C_VIEW,
          '💡 黄格子选订单号（下拉是【订单汇总】里的订单，也可以直接打），下面先是这张单的汇总，再是每一行（最多 200 行），右边按款式小计。'
          '单双成本＝这个款交货那个月的平均每双成本（【成本分摊表】分出来的），晚到的送货单录进去以后会变；没交的行、款式没登记的行没有成本。'
          '要改数据去【订单明细】改。', h2=36)
    home_link(ws, 'X1')
    ws.column_dimensions['X'].width = 10
    sel = OQ_SEL.split('!')[1].replace('$', '')          # C3
    first = next((x[OD_NO] for x in ctx.get(SH_ORD, []) if x.get(OD_NO) == '353651'), None) or \
        next((x[OD_NO] for x in ctx.get(SH_ORD, []) if x.get(OD_NO)), None)
    selector(ws, 'B3', '订单号', sel, first, f'={q(SH_OSUM)}!$B${OS_R0}:$B${OS_R1}', '@', '选订单号（【订单汇总】里的）')
    ws.merge_cells('C3:D3')
    S = OQ_SEL
    qs = odr(OD_QSEL)
    put(ws, 'E3', f'=IF(TRIM({S}&"")="","← 在黄格子选订单号",IF(COUNT({qs})=0,"✗ 【订单明细】里找不到这个订单号",'
                  f'"共 "&COUNT({qs})&" 行"&IF(COUNT({qs})>{OQ_N},"（只列前 {OQ_N} 行）","")))', F_KPI_L, align=AL, border=False)
    ws.merge_cells('E3:K3')
    ws.conditional_formatting.add('E3', FormulaRule(formula=['LEFT($E$3,1)="✗"'], font=F_RED))
    ws.row_dimensions[3].height = 26
    # 汇总（第 4 行标签、第 5 行数）
    sm_heads = [('B', '客户'), ('C', '下单日期'), ('D', '款式数'), ('E', '订单\n行数'), ('F', '订单双数'), ('G', '已交'),
                ('H', '未交'), ('I', '交货进度'), ('J', '交货收入'), ('K', '成本'), ('L', '毛利'), ('M', '毛利率'), ('N', '状态')]
    header(ws, 4, [('A', '汇总')] + sm_heads, H_DARK, height=30)
    first_ix = f'IFERROR(SMALL({qs},1),0)'
    vals = {
        'B': (f'=IF({first_ix}=0,"",INDEX({odr(OD_CUSTK)},{first_ix})&"")', None),
        'C': (f'=IF(COUNT(${OQ_ODATE}${OQ_R0}:${OQ_ODATE}${OQ_R1})=0,"",MIN(${OQ_ODATE}${OQ_R0}:${OQ_ODATE}${OQ_R1}))', DATE),
        'D': (f'=SUMIFS({odr(OD_OSFIRST)},{qs},">0")', INT),
        'E': (f'=COUNT({qs})', INT),
        'F': (f'=SUMIFS({odr(OD_OQ)},{qs},">0")', INT),
        'G': (f'=SUMIFS({odr(OD_DQ)},{qs},">0",{odr(OD_DQ)},">0")', INT),
        'H': ('=MAX(0,F5-G5)', INT),
        'I': ('=IF(F5>0,ROUND(G5/F5,4),"")', PCT),
        'J': (f'=ROUND(SUMIFS({odr(OD_AMT)},{qs},">0",{odr(OD_OK)},1),2)', MONEY),
        'K': (f'=ROUND(SUMIFS({odr(OD_COSTV)},{qs},">0"),2)', MONEY),
        'L': ('=ROUND(J5-K5,2)', MONEY),
        'M': ('=IF(J5<=0,"",ROUND(L5/J5,4))', PCT),
        'N': ('=IF(E5=0,"",IF(G5<=0,"未交",IF(G5<F5,"部分交","已交完")))', None),
    }
    put(ws, 'A5', '', F_TXTB, FILL_TOT)
    for c, (f, fmt) in vals.items():
        put(ws, f'{c}5', f, F_AUTOB, FILL_TOT, fmt, AR if fmt == MONEY else AC)
    ws.row_dimensions[5].height = 24
    _neg_red(ws, 'L5', 'L5')
    # 明细
    section(ws, 6, 'A', 'O', '逐行明细（【订单明细】里这张单的每一行）', H_VIEW)
    section(ws, 6, 'Q', 'W', '按款式小计', 'FF2F75B5')
    heads = [('A', '序号'), ('B', '款式编码'), ('C', '颜色及规格'), ('D', '码数'), ('E', '订单\n数量'), ('F', '交货日期'),
             ('G', '实交'), ('H', '结算\n单价'), ('I', '交货金额'), ('J', '单双\n成本'), ('K', '成本'), ('L', '毛利'), ('M', '毛利率'),
             ('N', '备注'), ('O', '校验')]
    header(ws, 7, heads, H_VIEW)
    header(ws, 7, [('Q', '款式'), ('R', '订单\n双数'), ('S', '已交'), ('T', '交货金额'), ('U', '成本'), ('V', '毛利'), ('W', '毛利率')],
           'FF2F75B5')
    for r in range(OQ_R0, OQ_R1 + 1):
        ix = f'${OQ_IDX}{r}'
        e = f'{ix}=0'
        g = lambda col: f'INDEX({odr(col)},{ix})'
        ws[f'{OQ_IDX}{r}'] = f'=IFERROR(SMALL({qs},ROW()-{OQ_R0 - 1}),0)'
        ws[f'{OQ_ODATE}{r}'] = f'=IF({e},"",IF(ISNUMBER({g(OD_DATE)}),INT({g(OD_DATE)}),""))'
        ws[f'{OQ_SF}{r}'] = f'=IF({e},"",IF({g(OD_OSFIRST)}=1,ROW()-{OQ_R0 - 1},""))'
        ws[f'A{r}'] = f'=IF({e},"",ROW()-{OQ_R0 - 1})'
        ws[f'B{r}'] = f'=IF({e},"",{g(OD_STY)}&"")'
        ws[f'C{r}'] = f'=IF({e},"",{g(OD_SPEC)}&"")'
        ws[f'D{r}'] = f'=IF({e},"",{g(OD_SIZE)}&"")'
        ws[f'E{r}'] = f'=IF({e},"",{g(OD_OQ)})'
        ws[f'F{r}'] = f'=IF({e},"",IF({g(OD_DDATE)}="","",{g(OD_DDATE)}))'
        ws[f'G{r}'] = f'=IF(OR({e},F{r}=""),"",{g(OD_DQ)})'
        ws[f'H{r}'] = f'=IF({e},"",{g(OD_PRICEU)})'
        ws[f'I{r}'] = f'=IF({e},"",{g(OD_AMTV)})'
        ws[f'J{r}'] = f'=IF({e},"",{g(OD_UCV)})'
        ws[f'K{r}'] = f'=IF({e},"",{g(OD_COSTV)})'
        ws[f'L{r}'] = f'=IF({e},"",{g(OD_GPV)})'
        ws[f'M{r}'] = f'=IF({e},"",IF(N(I{r})<=0,"",ROUND(N(L{r})/I{r},4)))'
        ws[f'N{r}'] = f'=IF({e},"",{g(OD_NOTE)}&"")'
        ws[f'O{r}'] = f'=IF({e},"",{g(OD_CHK)}&"")'
    cols = [CL(i) for i in range(1, 16)]
    style_rows(ws, OQ_R0, OQ_R1, cols, auto=cols,
               fmts={'E': INT, 'F': DATE, 'G': INT, 'H': MONEY, 'I': MONEY, 'J': MONEY, 'K': MONEY, 'L': MONEY, 'M': PCT},
               aligns={'B': AL, 'C': AL, 'N': AL, 'O': AL, 'H': AR, 'I': AR, 'J': AR, 'K': AR, 'L': AR})
    for r in range(OQ_R0, OQ_R1 + 1):
        for c in cols:
            ws[f'{c}{r}'].fill = FILL_NONE
    _chk_cf(ws, f'O{OQ_R0}:O{OQ_R1}', f'$O{OQ_R0}')
    _neg_red(ws, f'L{OQ_R0}:L{OQ_R1}', f'L{OQ_R0}')
    # 按款式小计
    sf = f'${OQ_SF}${OQ_R0}:${OQ_SF}${OQ_R1}'
    for k in range(OQ_SN):
        r = OQ_R0 + k
        ix = f'${OQ_SIDX}{r}'
        e = f'{ix}=0'
        ws[f'{OQ_SIDX}{r}'] = f'=IFERROR(SMALL({sf},ROW()-{OQ_R0 - 1}),0)'
        ws[f'Q{r}'] = f'=IF({e},"",INDEX($B${OQ_R0}:$B${OQ_R1},{ix}))'
        osk_sel = 'TRIM(' + S + '&"")&"|"&TRIM($Q' + str(r) + ')'
        key = f'{odr(OD_OSKEY)},{esc(osk_sel)}'
        ws[f'R{r}'] = f'=IF({e},"",SUMIFS({odr(OD_OQ)},{key}))'
        ws[f'S{r}'] = f'=IF({e},"",SUMIFS({odr(OD_DQ)},{key},{odr(OD_DQ)},">0"))'
        ws[f'T{r}'] = f'=IF({e},"",ROUND(SUMIFS({odr(OD_AMT)},{key},{odr(OD_OK)},1),2))'
        ws[f'U{r}'] = f'=IF({e},"",ROUND(SUMIFS({odr(OD_COSTV)},{key}),2))'
        ws[f'V{r}'] = f'=IF({e},"",ROUND(T{r}-U{r},2))'
        ws[f'W{r}'] = f'=IF({e},"",IF(N(T{r})<=0,"",ROUND(V{r}/T{r},4)))'
    scols = ['Q', 'R', 'S', 'T', 'U', 'V', 'W']
    style_rows(ws, OQ_R0, OQ_R0 + OQ_SN - 1, scols, auto=scols,
               fmts={'R': INT, 'S': INT, 'T': MONEY, 'U': MONEY, 'V': MONEY, 'W': PCT},
               aligns={'Q': AL, 'T': AR, 'U': AR, 'V': AR}, bold=['V'])
    for r in range(OQ_R0, OQ_R0 + OQ_SN):
        for c in scols:
            ws[f'{c}{r}'].fill = FILL_NONE
    _neg_red(ws, f'V{OQ_R0}:V{OQ_R0 + OQ_SN - 1}', f'V{OQ_R0}')
    _help_cols(ws, 7, [(OQ_IDX, '订单明细第几行'), (OQ_ODATE, '下单日期数字'), (OQ_SF, '款式首行★'), (OQ_SIDX, '第k个款')],
               OQ_R0, OQ_R1)
    ws.freeze_panes = f'C{OQ_R0}'
    print_setup(ws, '7:7', landscape=True)
    ws.print_area = f'A1:W{OQ_R1}'


def build(wb, ctx):
    build_sm(wb, ctx)
    build_alloc(wb, ctx)
    build_spl(wb, ctx)
    build_osum(wb, ctx)
    build_oq(wb, ctx)
