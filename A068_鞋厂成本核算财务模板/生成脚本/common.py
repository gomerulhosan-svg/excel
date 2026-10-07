# -*- coding: utf-8 -*-
"""A068 鞋厂成本核算财务模板 · 通用样式和小工具（跟 A064/A066 一个样：微软雅黑，彩色标题带 ＋ 💡提示）。
   各表的行列地址、表与表之间的「接口列」都在 layout.py。"""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL, column_index_from_string as CI
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink


def norm(x):
    """名称规范写法：去空格、半角括号改全角"""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({x})," ",""),"(","（"),")","）")'


def esc(x):
    """当 SUMIFS/COUNTIFS/MATCH 的条件用时，把名字里的 ~ * ? 转义成普通字符"""
    return f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"~","~~"),"*","~*"),"?","~?")'


# ─────────────────────────── 样式（跟 A064 一个样：微软雅黑，彩色标题带 ＋ 💡提示） ───────────────────────────
YH = '微软雅黑'
F_TITLE = Font(name=YH, sz=16, bold=True, color='FFFFFFFF')
F_TIP = Font(name=YH, sz=10, color='FFFF6600')
F_HDR = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_TXT = Font(name=YH, sz=10, color='FF000000')
F_TXTB = Font(name=YH, sz=10, bold=True, color='FF000000')
F_IN = Font(name=YH, sz=10, color='FF1F4E79')
F_AUTO = Font(name=YH, sz=10, color='FF404040')
F_AUTOB = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_NOTE = Font(name=YH, sz=9, color='FF808080')
F_HELP = Font(name=YH, sz=8, color='FFBFBFBF')
F_RED = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_KPI_L = Font(name=YH, sz=10, bold=True, color='FF1F3864')
F_KPI_V = Font(name=YH, sz=13, bold=True, color='FFC00000')
F_SEC = Font(name=YH, sz=11, bold=True, color='FFFFFFFF')
F_SEL = Font(name=YH, sz=11, bold=True, color='FF000000')


def fill(rgb):
    return PatternFill('solid', fgColor=rgb)


C_BASE, C_CASH, C_CONV, C_INV, C_RPT, C_AR, C_CHK, C_HOME = \
    'FF595959', 'FF2F75B5', 'FF7F7F7F', 'FF7030A0', 'FF375623', 'FFBF8F00', 'FF833C0C', 'FF1F3864'
C_IN, C_VAR, C_FIX, C_OTH, C_MISC = 'FFE2EFDA', 'FFDDEBF7', 'FFFFF2CC', 'FFF2F2F2', 'FFFFFFFF'   # 截图里各组的底色
C_IN_D, C_VAR_D, C_FIX_D = 'FF70AD47', 'FF5B9BD5', 'FFBF8F00'

FILL_TIP = fill('FFFFF2CC')
FILL_IN = fill('FFFFF7E0')      # 手工录入：淡黄
FILL_PASTE = fill('FFEAF1FB')   # 粘贴区：淡蓝
FILL_AUTO = fill('FFF2F2F2')    # 自动：淡灰
FILL_SUB = fill('FFDDEBF7')     # 小计
FILL_TOT = fill('FFFCE4D6')     # 合计 / 利润
FILL_SEL = fill('FFFFFF00')     # 选择格（亮黄）
FILL_OK = fill('FFE2EFDA')
FILL_WARN = fill('FFFFC7CE')
FILL_NONE = PatternFill(fill_type=None)

thin = Side(style='thin', color='FFBFBFBF')
BD = Border(left=thin, right=thin, top=thin, bottom=thin)
NOBD = Border()
AC = Alignment(horizontal='center', vertical='center')
ACW = Alignment(horizontal='center', vertical='center', wrap_text=True)
AL = Alignment(horizontal='left', vertical='center')
ALW = Alignment(horizontal='left', vertical='center', wrap_text=True)
AR = Alignment(horizontal='right', vertical='center')

MONEY = '#,##0.00;[Red]-#,##0.00;"-"'
MONEY0 = '#,##0;[Red]-#,##0;"-"'
DATE = 'yyyy/mm/dd'
DTIME = 'yyyy/mm/dd hh:mm'
INT = '0;-0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
MONTH = 'm"月"'


def put(ws, coord, value=None, font=None, fill_=None, fmt=None, align=None, border=True):
    c = ws[coord]
    if value is not None:
        c.value = value
    if font is not None:
        c.font = font
    if fill_ is not None:
        c.fill = fill_
    if fmt is not None:
        c.number_format = fmt
    if align is not None:
        c.alignment = align
    if border:
        c.border = BD
    return c


def title(ws, text, last_col, color, tip=None, h1=33, h2=None):
    """第 1 行标题带 ＋ 第 2 行 💡 提示；提示长就把第 2 行加高"""
    if h2 is None:
        width = sum((ws.column_dimensions[CL(i)].width or 9) for i in range(1, CI(last_col) + 1))
        per_line = max(20, int(width / 1.9))
        lines = -(-len(tip or '') // per_line)
        h2 = max(24, 16 * lines + 6)
    ws.merge_cells(f'A1:{last_col}1')
    put(ws, 'A1', text, F_TITLE, fill(color), align=AC, border=False)
    ws.row_dimensions[1].height = h1
    if tip:
        ws.merge_cells(f'A2:{last_col}2')
        put(ws, 'A2', tip, F_TIP, FILL_TIP, align=ALW, border=False)
        ws.row_dimensions[2].height = h2
    ws.sheet_properties.tabColor = color[2:] if len(color) == 8 else color
    ws.sheet_view.showGridLines = False


def header(ws, row, cols_texts, color, font=F_HDR, height=34):
    for col, t in cols_texts:
        put(ws, f'{col}{row}', t, font, fill(color), align=ACW)
    ws.row_dimensions[row].height = height


def section(ws, row, c1, c2, text, color):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    put(ws, f'{c1}{row}', text, F_SEC, fill(color), align=AL)
    for i in range(CI(c1) + 1, CI(c2) + 1):
        ws.cell(row=row, column=i).border = BD


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def hide(ws, *cols):
    for c in cols:
        ws.column_dimensions[c].hidden = True


def link(cell, sheet, ref='A1'):
    """表内跳转：写成 location，WPS / Excel 都能点过去"""
    v = cell.value
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{sheet}'!{ref}",
                               display=None if v is None or str(v).startswith('=') else str(v))
    return cell


def style_rows(ws, r0, r1, cols, auto=(), fmts=None, aligns=None, fills=None, bold=()):
    """数据区：白底细灰框；公式列淡灰；fills 可给手工列指定底色"""
    fmts, aligns, fills = fmts or {}, aligns or {}, fills or {}
    for r in range(r0, r1 + 1):
        for col in cols:
            c = ws[f'{col}{r}']
            is_auto = col in auto
            c.font = F_AUTOB if col in bold else (F_AUTO if is_auto else F_IN)
            c.fill = FILL_AUTO if is_auto else fills.get(col, FILL_NONE)
            c.border = BD
            c.alignment = aligns.get(col, AC)
            if col in fmts:
                c.number_format = fmts[col]


def dv_list(ws, sqref, formula, prompt=None, stop=True, blank=True):
    dv = DataValidation(type='list', formula1=formula, allow_blank=blank, showErrorMessage=stop,
                        errorStyle='stop' if stop else 'warning')
    if prompt:
        dv.promptTitle, dv.prompt = '提示', prompt
        dv.showInputMessage = True
    dv.errorTitle = '不在清单里'
    dv.error = '请从下拉里选；清单里没有的，先到【基础资料】/【往来单位】/【款式档案】里登记'
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def dv_date(ws, sqref):
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050', allow_blank=True,
                        showErrorMessage=True, errorStyle='warning', errorTitle='日期', error='要像 2026/9/1 这样填日期')
    ws.add_data_validation(dv)
    dv.add(sqref)
    return dv


def selector(ws, cell_lbl, lbl, cell_in, value, dv_formula=None, fmt=None, prompt=None):
    """报表顶上的选择格：标签 ＋ 亮黄输入格"""
    put(ws, cell_lbl, lbl, F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, cell_in, value, F_SEL, FILL_SEL, fmt=fmt, align=AC)
    if dv_formula:
        dv = DataValidation(type='list', formula1=dv_formula, allow_blank=True, showErrorMessage=False)
        if prompt:
            dv.promptTitle, dv.prompt, dv.showInputMessage = '提示', prompt, True
        ws.add_data_validation(dv)
        dv.add(cell_in.replace('$', ''))


def date_parse(x):
    """把粘贴进来的各种日期写法变成真日期（取日期部分）：
       真日期/日期时间 → 取整；"2026-09-01 08:24:04" / "2026/9/1" / "2026.9.1" / "2026年9月1日" / "20260901" 文本 → 拆年月日。
       只有纯 8 位数字才按 yyyymmdd 拆（"2026-9-1" 也是 8 个字，不能按位置拆）；小于 2000 年的数（比如只有时间）算看不懂"""
    s = (f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({x}),"/","-"),".","-"),"年","-"),"月","-"),"日"," ")')
    dp = f'LEFT({s},FIND(" ",{s}&" ")-1)'
    return (f'IF({x}="","",IF(ISNUMBER({x}),IF({x}>19000000,DATE(INT({x}/10000),MOD(INT({x}/100),100),MOD({x},100)),'
            f'IF({x}<36526,"",INT({x}))),'
            f'IFERROR(IF(ISERROR(FIND("-",{dp})),DATE(LEFT({dp},4),MID({dp},5,2),RIGHT({dp},2)),'
            f'DATE(LEFT({dp},4),MID({dp},6,FIND("-",{dp},6)-6),MID({dp},FIND("-",{dp},6)+1,2))),"")))')


def num(x):
    """粘贴来的金额可能是文本、带 ¥ 或千分位逗号"""
    return f'IF({x}="",0,IFERROR(--SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({x},"¥",""),"￥",""),",",""),0))'


# ─────────────────────────── 查看表共用：回首页链接、按条件列清单 ───────────────────────────
def home_link(ws, coord, home='首页'):
    c = put(ws, coord, '← 回首页', Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'),
            fill('FFFFFFFF'), align=AC, border=False)
    link(c, home, 'A1')
    return c


def counter(ws, col, r0, n, cond, block=50):
    """隐藏的计数列：第 r0 行起 n 行，第 i 行＝前 i 行里满足条件的个数（cond(i) 给出第 i 行的条件公式）。
       条件先放在另一列（自动找本表右边空列），计数按 50 行一段求和：公式链只有几十层深（不会太深算不动），改一格也不用整列重算。
       取第 k 个满足条件的行：kth(k, ...)。"""
    idx = getattr(ws, '_ccn', CI('CA'))
    ws._ccn = idx + 1
    cc = CL(idx)
    for i in range(n):
        r = r0 + i
        ws[f'{cc}{r}'] = f'=IF({cond(i)},1,0)'
        bs = r0 + (i // block) * block
        base = f'{col}{bs - 1}+' if bs > r0 else ''
        ws[f'{col}{r}'] = f'={base}SUM({cc}{bs}:{cc}{r})'
        ws[f'{col}{r}'].font = ws[f'{cc}{r}'].font = F_HELP
    ws.column_dimensions[cc].hidden = True


def kth(k, col, r0, n):
    """第 k 个满足条件的是第几行（1 起；没有就 0）"""
    last = f'${col}${r0 + n - 1}'
    return f'IF({k}>{last},0,IFERROR(MATCH({k},${col}${r0}:{last},0),0))'


def cnt(col, r0, n):
    return f'${col}${r0 + n - 1}'


def skey(ws, col, r0, n, cond, datef):
    """按日期排序的清单用：满足条件的行放「日期×10000＋第几行」，不满足放空；返回放个数的格子"""
    for i in range(n):
        r = r0 + i
        ws[f'{col}{r}'] = f'=IF({cond(i)},{datef(i)}*10000+{i + 1},"")'
        ws[f'{col}{r}'].font = F_HELP
    c = f'{col}{r0 - 1}' if r0 > 1 else f'{col}{r0 + n}'
    ws[c] = f'=COUNT({col}{r0}:{col}{r0 + n - 1})'
    ws[c].font = F_HELP
    return f'${c[0:len(c.rstrip("0123456789"))]}${c[len(c.rstrip("0123456789")):]}'


def ksorted(k, col, r0, n, ncell):
    """按日期排第 k 笔是源表第几行（1 起；没有就 0）"""
    return f'IF({k}>{ncell},0,MOD(SMALL(${col}${r0}:${col}${r0 + n - 1},{k}),10000))'


def print_setup(ws, rows=None, landscape=True, fit_width=True):
    """打印：横向、一页宽、每页重复表头"""
    ws.page_setup.orientation = 'landscape' if landscape else 'portrait'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    if fit_width:
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
    if rows:
        ws.print_title_rows = rows
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.5


def ym_month(date_cell, year_cell):
    """日期 → 本年第几月（1～12）；不是日期、不在本年度 → 0"""
    return f'IF(ISNUMBER({date_cell}),IF(YEAR({date_cell})={year_cell},MONTH({date_cell}),0),0)'


def month_end(year_cell, m):
    """本年 m 月最后一天"""
    return f'DATE({year_cell},{m}+1,0)'


def txt(x):
    """科目编码统一成文本（数字 1002 和文字 "1002" 当成一个）"""
    return f'TRIM({x}&"")'
