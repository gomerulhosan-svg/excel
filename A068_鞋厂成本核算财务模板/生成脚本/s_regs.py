# -*- coding: utf-8 -*-
"""录入用的四张登记表：【送货单登记】【外发加工登记】【订单明细】【工资登记】（第一阶段）。
   每张：第 1 行标题、第 2 行 💡提示、第 3 行汇总、第 4 行表头、第 5 行起数据；右边隐藏的是 layout.py 里的 ★接口列，
   再往右是本表自己用的帮手列（也隐藏）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from common import *
from layout import *

YR, OPEN, LAGD, CUST = (P[k] for k in ('YEAR', 'OPEN', 'LAGD', 'CUST'))
ST = f'MAX({OPEN},DATE({YR},1,1))'                                       # 本账从哪天起（建账日期；在上年就从 1/1 起）
OM = f'IF(ISNUMBER({OPEN}),IF(YEAR({OPEN})={YR},MONTH({OPEN}),1),1)'     # 建账月：建账日期在本年度＝它的月份，否则 1
H_IN, H_AUTO = C_IN, 'FF548235'                                          # 表头：录入列蓝、自动列绿
F_GOOD = Font(name=YH, sz=10, bold=True, color='FF00B050')
F_WAIT = Font(name=YH, sz=10, color='FF7F7F7F')
QTYF = '#,##0.###;[Red]-#,##0.###'
INTR = '0;[Red]-0;0'
SM_UNIT2D = f"{q(SH_SM)}!${sm_col('单双', 1)}${SM_R0}:${sm_col('单双', 12)}${SM_R1}"   # _款式月「单双」块：款式行 × 12 个月


# ─────────────────────────── 小工具 ───────────────────────────
def _blank(x):
    return f'TRIM({x}&"")=""'


def _bad(x):
    """填了、但不是数字"""
    return f'AND(TRIM({x}&"")<>"",NOT(ISNUMBER({x})))'


def _txtnum(x):
    """填的是「文本格式的数字」（网银 / 别的表粘贴来的 "12,000.00"、" 50 "、"¥50"：看着是数字，进不了账）"""
    t = f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE(TRIM({x}&""),",",""),"¥",""),"￥","")'
    return f'AND(TRIM({x}&"")<>"",NOT(ISNUMBER({x})),ISNUMBER(--{t}))'


def _txtmsg(what):
    return f'"✗ {what}是文本格式的数字（粘贴来的，看着是数字、算不进账）：选中这几列 → 数据 → 分列 → 完成；或点格子左上角绿色小三角 → 转换为数字"'


def _s(text):
    return f'"{text}"'


def _chain(conds, default):
    """校验：按顺序，命中第一个就停"""
    f = default
    for c, msg in reversed(conds):
        f = f'IF({c},{msg},{f})'
    return f


def _kpi(ws, lbl, label, val, formula, fmt=MONEY):
    """第 3 行汇总：标签格 ＋ 数值格（可以是合并区域 'A3:B3'）"""
    for rg in (lbl, val):
        if ':' in rg:
            ws.merge_cells(rg)
    put(ws, lbl.split(':')[0], label, F_KPI_L, fill('FFD9E1F2'), align=ACW)
    put(ws, val.split(':')[0], formula, F_KPI_V, fill('FFD9E1F2'), fmt, AC)


def _heads(ws, hdr, heads, autos):
    for col, t in heads:
        put(ws, f'{col}{hdr}', t, F_HDR, fill(H_AUTO if col in autos else H_IN), align=ACW)
    ws.row_dimensions[hdr].height = 46


def _hidden(ws, hdr, r0, r1, labels, fmts=None):
    """隐藏列：表头写个小字说明、整列小灰字、隐藏"""
    fmts = fmts or {}
    for c, t in labels:
        ws[f'{c}{hdr}'] = t
        ws[f'{c}{hdr}'].font = F_HELP
        for r in range(r0, r1 + 1):
            ws[f'{c}{r}'].font = F_HELP
            if c in fmts:
                ws[f'{c}{r}'].number_format = fmts[c]
    hide(ws, *[c for c, _ in labels])


def _chk_cf(ws, col, r0, r1):
    rg = f'{col}{r0}:{col}{r1}'
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="✗"'], fill=FILL_WARN, font=F_RED))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="⚠"'], fill=fill('FFFFEB9C')))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="√"'], font=F_GOOD))
    ws.conditional_formatting.add(rg, FormulaRule(formula=[f'LEFT(${col}{r0},1)="⏳"'], font=F_WAIT))


def _put_rows(ws, rows, r0):
    for i, row in enumerate(rows):
        for c, v in row.items():
            if v is not None:
                ws[f'{c}{r0 + i}'] = v


def _dv_whole(ws, sqref, lo, hi, msg):
    dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, errorStyle='warning', errorTitle='数字', error=msg)
    ws.add_data_validation(dv)
    dv.add(sqref)


# ─────────────────────────── 送货单 / 外发 共用：金额、成本月、记账日期、成本键…… ───────────────────────────
def _close(sm):
    """送货月 sm（1～12）那个月的结账日期（【基础资料】⑩；空＝没结账）。只在 sm 一定是 1～12 的分支里用（INDEX 的 k 不能是 0）"""
    return f'INDEX({CLOSE_DATES},{sm})'


def _srow(sty):
    """「用在哪」在【款式档案】第几个（1 起；空着或没登记 0）——本表帮手列用"""
    t = f'TRIM({sty}&"")'
    return f'=IF({_blank(sty)},0,IFERROR(MATCH({esc(t)},{ST_KEYS},0),0))'


def _atype(k, cm):
    """款 / 公：登记的款式（第 k 个）→ 款；否则 公。停产在成本引擎里处理（_款式月 停X），这里不看，免得改订单时整张表变脏"""
    return f'IF({k}=0,"公","款")'


def _bill_row(ws, r, c, hdr, comp):
    """c：角色 → 列字母。写 SEQ AMT SM RM LAG CM PDATE SROW ATYPE CKEY UNITK UROW OK LATEK CMV LAGV 这些公式"""
    g = {k: f'{v}{r}' for k, v in c.items()}
    A, B, C, D = g['SEQ'], g['DATE'], g['RDATE'], g['SUP']
    SM, RM, CM, OK = g['SM'], g['RM'], g['CM'], g['OK']
    i = f'ROW()-{hdr}'
    ws[A] = f'=IF(COUNTA({c["DATE"]}{r}:{c["NOTE"]}{r})=0,"",{i})'
    ws[g['AMT']] = (f'=IF(ISNUMBER({g["AMTIN"]}),ROUND({g["AMTIN"]},2),'
                    f'IF({_blank(g["AMTIN"])},ROUND(N({g["QTY"]})*N({g["PRICE"]}),2),0))')
    ws[SM] = f'=IF({A}="",0,IF(ISNUMBER({B}),IF(YEAR({B})<{YR},0,IF(YEAR({B})>{YR},13,MONTH({B}))),-1))'
    ws[RM] = (f'=IF({A}="",0,IF({_blank(C)},{SM},'
              f'IF(ISNUMBER({C}),IF(YEAR({C})<{YR},0,IF(YEAR({C})>{YR},13,MONTH({C}))),-1)))')
    ws[g['LAG']] = f'=IF(AND(ISNUMBER({B}),ISNUMBER({C})),INT({C})-INT({B}),0)'
    # 成本月：下一年度/不是日期 0；建账前送的（含上年）→ 收单也在建账前 0，否则 MAX(收单月,建账月)；
    #   否则（这时 SM 一定是 1～12）：送货月填了结账日期、且 收单日期（空＝送货日期）晚于结账日期 → 收单月（次年 → 0 ✗）；否则送货月
    cd = _close(SM)
    eff = f'INT(IF(ISNUMBER({C}),{C},{B}))'
    ws[CM] = (f'=IF({A}="",0,IF(OR({SM}<0,{SM}=13),0,IF({B}<{ST},'
              f'IF(ISNUMBER({C}),IF({C}<{ST},0,IF({RM}=13,0,MAX({RM},{OM}))),0),'
              f'IF(ISNUMBER({cd}),IF({eff}>{cd},IF({RM}=13,0,{RM}),{SM}),{SM}))))')
    ws[g['PDATE']] = (f'=IF({CM}<1,"",IF(AND({CM}={SM},{B}>={ST}),INT({B}),'
                      f'IF(IF(ISNUMBER({C}),AND({RM}={CM},{C}>={ST}),FALSE),INT({C}),MAX(DATE({YR},{CM},1),{ST}))))')
    sty = g['STY']
    ws[g['SROW']] = _srow(sty)
    ws[g['ATYPE']] = f'=IF({A}="","",{_atype(g["SROW"], CM)})'
    ws[g['CKEY']] = f'=IF({OK}=1,IF({g["ATYPE"]}="款","款|"&TRIM({sty}&"")&"|{comp}","公||{comp}"),"")'
    ws[g['UNITK']] = f'=IF({A}="","",TRIM({D}&""))'
    ws[g['UROW']] = f'=IF({g["UNITK"]}="",0,IFERROR(MATCH({esc(g["UNITK"])},{UN_NAMES},0),0))'
    ws[OK] = f'=IF(AND({g["CHK"]}<>"",LEFT({g["CHK"]},1)<>"✗"),1,0)'
    ws[g['LATEK']] = f'=IF({OK}=1,IF({g["LAG"]}>={LAGD},{g["LAG"]}*100000+{i},""),"")'
    ws[g['CMV']] = f'=IF({CM}>=1,{CM},"")'
    ws[g['LAGV']] = f'=IF(AND(ISNUMBER({B}),ISNUMBER({C})),{g["LAG"]},"")'


def _bill_checks(r, c, t):
    """送货单 / 外发 的校验条件（按顺序）。t：文案（dname 日期叫法、who 单位叫法、want 应有的单位类型、what 送货/加工、pre 建账前提醒用）"""
    g = {k: f'{v}{r}' for k, v in c.items()}
    B, C = g['DATE'], g['RDATE']
    SM, CM, AMT = g['SM'], g['CM'], g['AMT']
    pre = f'{B}<{ST}'                                                    # 只在 B 是日期以后才用
    c_pre = f'IF(ISNUMBER({C}),{C}<{ST},{_blank(C)})'
    q_, p_, a_ = g['QTY'], g['PRICE'], g['AMTIN']
    diff = f'IF(AND(ISNUMBER({q_}),ISNUMBER({p_}),ISNUMBER({a_})),ABS(ROUND({a_}-{q_}*{p_},2)),0)'
    cd = _close(SM)                         # 下面用到它的几条都在「建账前」那条之后：SM 一定是 1～12
    cdt = f'TEXT({cd},"m/d")'
    closed_blank = f'IF({SM}>=1,IF({SM}<=12,IF(ISNUMBER({cd}),{_blank(C)},FALSE),FALSE),FALSE)'
    k = g['SROW']
    return [
        (f'NOT(ISNUMBER({B}))', _s(f'✗ {t["dname"]}要填成日期（如 2026/9/5）')),
        (f'{SM}=13', _s('✗ 下一年度的单，这本账不记（记到下一年的账里）')),
        (f'AND({pre},{c_pre})', _s(f'✗ {t["what"]}、收单都在建账以前：这笔欠款放【往来单位】期初应付，这里不记')),
        (_blank(g['SUP']), _s(f'✗ 没选{t["who"]}')),
        (f'{g["UROW"]}=0', _s(f'✗ {t["who"]}不在【往来单位】：先去登记（名称要一模一样）')),
        (_bad(C), _s('✗ 收单日期要填成日期（不知道就空着）')),
        (f'IF(ISNUMBER({C}),INT({C})<INT({B}),FALSE)', _s(f'✗ 收单日期比{t["dname"]}还早：看看是不是填反了')),
        (f'{CM}<1', _s('✗ 收单日期在下一年度：这本账不记（记到下一年的账里）')),
        (f'OR({_txtnum(q_)},{_txtnum(p_)},{_txtnum(a_)})', _txtmsg('数量 / 单价 / 金额')),
        (f'OR({_bad(q_)},{_bad(p_)},{_bad(a_)})', _s('✗ 数量、单价、金额要填数字（不要带单位、文字）')),
        (f'{AMT}=0', _s('✗ 没有金额（数量×单价是 0，金额也没填）')),
        (pre, f'"⚠ 建账前{t["pre"]}，单子后来才到：算到 "&{CM}&" 月"'),
        # 到这里 B 不早于建账日期：CM≠SM 只可能是「送货月已结账、结账日期以后才收到」挪到了收单月（这时 C 一定是日期）
        (f'{CM}<>{SM}', f'"⚠ "&{SM}&" 月已结账（结账日期 "&{cdt}&"），这张单 "&TEXT({C},"m/d")&" 才收到：算到收单月 "&{CM}&" 月"'),
        (closed_blank, f'"⚠ "&{SM}&" 月已结账（"&{cdt}&"）：要填收单日期（空着就算回 "&{SM}&" 月；结账以后才录的单会改动已结账的数）"'),
        (_blank(C), _s(f'⚠ 没填收单日期（按{t["dname"]}算；填上才看得出晚到多久）')),
        (f'IF({g["UROW"]}=0,FALSE,INDEX({UN_TYPES_R},{g["UROW"]})&""<>"{t["want"]}")',
         _s(f'⚠ {t["who"]}类型不是「{t["want"]}」（去【往来单位】看看类型）')),
        (f'{diff}>1', f'"⚠ 金额跟 数量×单价 差 "&TEXT({diff},"0.00")&" 元（照单子填的就不用改）"'),
        (f'AND(NOT({_blank(g["STY"])}),{g["ATYPE"]}="公")',
         '"⚠ 用在哪的款式没在【款式档案】：先按公共分（登记了自动改记到这个款）"'),
    ]


# ═══════════════════════════ 送货单登记 ═══════════════════════════
DN_UROW, DN_MROW, DN_SROW = 'AI', 'AJ', 'AK'   # 本表帮手列：供应商在【往来单位】第几行；品名在【品名档案】第几行；用在哪在【款式档案】第几个（0＝没登记）


def build_dn(wb, ctx):
    ws = wb[SH_DN]
    hdr, r0, r1 = DN_HDR, DN_R0, DN_R1
    widths(ws, {'A': 5, 'B': 11, 'C': 11, 'D': 16, 'E': 10, 'F': 16, 'G': 8, 'H': 5, 'I': 8, 'J': 8, 'K': 11, 'L': 10,
                'M': 20, 'N': 10, 'O': 7, 'P': 7, 'Q': 11, 'R': 40, 'S': 10})
    title(ws, '送 货 单 登 记（供应商送货单 · 一行一个品名 · 欠供应商的材料款）', DN_CHK, C_IN,
          '💡 照供应商的送货单抄：一个品名一行，单号、日期每行都写。「送货日期」写单子上的日期（材料成本算到那个月）；'
          '「收单日期」写拿到单子那天（供应商常常隔几个月才拿单来，填上就看得出晚到多久）。'
          '金额可以不填（＝数量×单价）；单子上写的金额跟 数量×单价 对不上（四舍五入）就照单子填金额；退货数量填负数。'
          '「用在哪」只有确定是某一个款专用的才填款式编码（比如防水台大底只给 24D16 用），不确定就空着（按交货双数分到各款）；'
          '款式在【款式档案】改了「停产」的，它最后一次交货以后的单按公共分。'
          '供应商要先在【往来单位】登记；品名最好在【品名档案】登记类别（没登记的算「未分类」）。'
          '某个月在【基础资料】⑩ 填了结账日期：那个月送的货、结账日期以后才收到的单，算到收单那个月（结了账的月份不再变）；'
          '所以结账以后补来的单一定要填收单日期，空着就按送货日期算回已结账的月份。'
          '记错了就改或清空那几格，不要插行删行（右边有隐藏公式）。付材料款在【资金日记账】记。')
    c = dict(SEQ=DN_SEQ, DATE=DN_DATE, RDATE=DN_RDATE, SUP=DN_SUP, NO=DN_NO, QTY=DN_QTY, PRICE=DN_PRICE, AMTIN=DN_AMTIN,
             STY=DN_STY, NOTE=DN_NOTE, CMV=DN_CMV, LAGV=DN_LAGV, CHK=DN_CHK, AMT=DN_AMT, CM=DN_CM, PDATE=DN_PDATE,
             ATYPE=DN_ATYPE, CKEY=DN_CKEY, UNITK=DN_UNITK, SM=DN_SM, RM=DN_RM, LAG=DN_LAG, OK=DN_OK, LATEK=DN_LATEK,
             UROW=DN_UROW, SROW=DN_SROW)
    t = dict(dname='送货日期', who='供应商', want='材料供应商', what='送货', pre='送的货')
    _put_rows(ws, ctx.get(SH_DN, []), r0)
    namek_r, notek_r, amt_r = dnr(DN_NAMEK), dnr(DN_NOTEK), dnr(DN_AMT)
    mt_names = rng(SH_MAT, MT_NAME, MT_R0, MT_R1)
    for r in range(r0, r1 + 1):
        _bill_row(ws, r, c, hdr, '材料')
        g = lambda col: f'{col}{r}'
        A, i = g(DN_SEQ), f'ROW()-{hdr}'
        nk = g(DN_NAMEK)
        ws[nk] = f'=IF({_blank(g(DN_NAME))},"",{norm(g(DN_NAME))})'
        ws[g(DN_MROW)] = f'=IF({nk}="",0,IFERROR(MATCH({esc(nk)},{MT_KEYS},0),0))'
        mrow = g(DN_MROW)
        ws[g(DN_CAT)] = (f'=IF({A}="","",IF({mrow}=0,"{MC_NONE}",'
                         f'IF(INDEX({MT_CATS},{mrow})&""="","{MC_NONE}",INDEX({MT_CATS},{mrow})&"")))')
        ws[g(DN_NFIRST)] = (f'=IF(AND({g(DN_OK)}=1,{nk}<>""),IF(COUNTIFS(${DN_NAMEK}${hdr}:{DN_NAMEK}{r - 1},{esc(nk)},'
                            f'${DN_OK}${hdr}:{DN_OK}{r - 1},1)=0,{i},""),"")')
        ws[g(DN_NNEW)] = f'=IF({g(DN_NFIRST)}="","",IF({mrow}=0,{g(DN_NFIRST)},""))'
        ws[g(DN_NOTEK)] = f'=IF({_blank(g(DN_NO))},"",{g(DN_UNITK)}&"|"&TRIM({g(DN_NO)}&""))'
        nt = g(DN_NOTEK)
        ws[g(DN_DOCT)] = (f'=IF({nt}="","",IF(IFERROR(MATCH({esc(nt)},{notek_r},0),0)={i},'
                          f'SUMIFS({amt_r},{notek_r},{esc(nt)}),""))')
        ws[g(DN_CHK)] = f'=IF({A}="","",{_chain(_bill_checks(r, c, t), _s("√"))})'
    cols = [DN_SEQ, DN_DATE, DN_RDATE, DN_SUP, DN_NO, DN_NAME, DN_COLOR, DN_UNIT, DN_QTY, DN_PRICE, DN_AMTIN, DN_STY, DN_NOTE,
            DN_CAT, DN_CMV, DN_LAGV, DN_DOCT, DN_CHK]
    autos = [DN_SEQ, DN_CAT, DN_CMV, DN_LAGV, DN_DOCT, DN_CHK]
    style_rows(ws, r0, r1, cols, auto=autos,
               fmts={DN_DATE: DATE, DN_RDATE: DATE, DN_QTY: QTYF, DN_PRICE: MONEY, DN_AMTIN: MONEY, DN_CMV: '0"月"',
                     DN_LAGV: INTR, DN_DOCT: MONEY, DN_NO: '@'},
               aligns={DN_SUP: AL, DN_NAME: AL, DN_NOTE: AL, DN_CHK: AL, DN_PRICE: AR, DN_AMTIN: AR, DN_DOCT: AR, DN_QTY: AR},
               fills={col: FILL_IN for col in cols if col not in autos})
    heads = [(DN_SEQ, '序号'), (DN_DATE, '送货日期\n（单子上的）'), (DN_RDATE, '收单日期\n（拿到单子）'), (DN_SUP, '供应商'),
             (DN_NO, '单号'), (DN_NAME, '品名及规格'), (DN_COLOR, '颜色'), (DN_UNIT, '单位'), (DN_QTY, '数量'), (DN_PRICE, '单价'),
             (DN_AMTIN, '金额\n（可空）'), (DN_STY, '用在哪\n（款式，可空）'), (DN_NOTE, '备注'), (DN_CAT, '材料类别'),
             (DN_CMV, '成本\n月份'), (DN_LAGV, '滞后\n天数'), (DN_DOCT, '单据合计'), (DN_CHK, '校验')]
    _heads(ws, hdr, heads, autos)
    _hidden(ws, hdr, r0, r1, [(DN_AMT, '金额'), (DN_CM, '成本月'), (DN_PDATE, '记账日期'), (DN_ATYPE, '款/公'), (DN_CKEY, '成本键'),
                               (DN_UNITK, '供应商'), (DN_SM, '送货月'), (DN_RM, '收单月'), (DN_LAG, '滞后天数'),
                               (DN_NAMEK, '品名规范写法'), (DN_NFIRST, '品名首次★'), (DN_NNEW, '没登记品名★'),
                               (DN_NOTEK, '供应商|单号'), (DN_OK, '能记账'), (DN_LATEK, '晚到键★'),
                               (DN_UROW, '往来单位第几行'), (DN_MROW, '品名档案第几行'), (DN_SROW, '款式档案第几个')],
            fmts={DN_AMT: MONEY, DN_PDATE: DATE})
    # 第 3 行汇总
    ok, chk = dnr(DN_OK), dnr(DN_CHK)
    _kpi(ws, 'A3:B3', '行数', 'C3', f'=COUNT({dnr(DN_SEQ)})', INT)
    _kpi(ws, 'D3:E3', '本年金额\n（能记账的）', 'F3:G3', f'=SUMIFS({amt_r},{ok},1)')
    _kpi(ws, 'H3:I3', '其中直接\n记到款式', 'J3:K3', f'=SUMIFS({amt_r},{ok},1,{dnr(DN_ATYPE)},"款")')
    _kpi(ws, 'L3', '✗ 待改', 'M3', f'=COUNTIF({chk},"✗*")', INT)
    _kpi(ws, 'N3', '⚠ 提醒', 'O3', f'=COUNTIF({chk},"⚠*")', INT)
    _kpi(ws, 'P3:Q3', f'="晚到（≥"&{LAGD}&"天）"', 'R3', f'=COUNT({dnr(DN_LATEK)})&" 行"', INT)
    ws.row_dimensions[3].height = 30
    _chk_cf(ws, DN_CHK, r0, r1)
    ws.conditional_formatting.add(f'{DN_LAGV}{r0}:{DN_LAGV}{r1}', FormulaRule(
        formula=[f'${DN_LATEK}{r0}<>""'], fill=fill('FFFCE4D6'), font=F_RED))
    dv_date(ws, f'{DN_DATE}{r0}:{DN_RDATE}{r1}')
    dv_list(ws, f'{DN_SUP}{r0}:{DN_SUP}{r1}', f'={UN_NAMES}', '【往来单位】登记过的供应商')
    dv_list(ws, f'{DN_NAME}{r0}:{DN_NAME}{r1}', f'={mt_names}', '【品名档案】里的品名；没有的直接写', stop=False)
    dv_list(ws, f'{DN_STY}{r0}:{DN_STY}{r1}', f'={ST_CODES}', '确定只用在这个款才填；不确定空着（按交货双数分）', stop=False)
    ws.auto_filter.ref = f'A{hdr}:{DN_CHK}{r1}'
    ws.freeze_panes = f'E{r0}'


# ═══════════════════════════ 外发加工登记 ═══════════════════════════
OT_UROW, OT_SROW = 'AC', 'AD'           # 本表帮手列：加工厂在【往来单位】第几行；用在哪在【款式档案】第几个（0＝没登记）


def build_ot(wb, ctx):
    ws = wb[SH_OUT]
    hdr, r0, r1 = OT_HDR, OT_R0, OT_R1
    widths(ws, {'A': 5, 'B': 11, 'C': 11, 'D': 16, 'E': 10, 'F': 11, 'G': 10, 'H': 10, 'I': 8, 'J': 5, 'K': 8, 'L': 11,
                'M': 22, 'N': 7, 'O': 7, 'P': 40, 'Q': 10})
    title(ws, '外 发 加 工 登 记（外发加工厂的加工单 · 欠加工厂的加工费）', OT_CHK, C_IN,
          '💡 外发出去加工（鞋面针车、印花、绣花……）的加工费在这里记，照加工厂的回货单 / 结算单一行一行抄；'
          '材料本身已经在【送货单登记】算了，这里只记加工费，不要重复。一般按款式外发，「用在哪」填款式编码，加工费直接记到这个款；'
          '没法分的（比如一批印花几个款一起）空着，按交货双数分到各款。「单据日期」写单子上的日期（成本算到那个月），'
          '「收单日期」写拿到单子那天。金额可以不填（＝数量×单价）。订单号可以不填，只是记一下。'
          '跟送货单一样：单据月份在【基础资料】⑩ 填了结账日期的，结账日期以后才收到的单算到收单月（结账以后补来的一定要填收单日期）；'
          '款式「停产」了，它最后一次交货以后的加工费按公共分。'
          '加工厂要先在【往来单位】登记（类型选「外发加工厂」）。付加工费在【资金日记账】记「付加工费」。不要插行删行（右边有隐藏公式）。')
    c = dict(SEQ=OT_SEQ, DATE=OT_DATE, RDATE=OT_RDATE, SUP=OT_SUP, NO=OT_NO, QTY=OT_QTY, PRICE=OT_PRICE, AMTIN=OT_AMTIN,
             STY=OT_STY, NOTE=OT_NOTE, CMV=OT_CMV, LAGV=OT_LAGV, CHK=OT_CHK, AMT=OT_AMT, CM=OT_CM, PDATE=OT_PDATE,
             ATYPE=OT_ATYPE, CKEY=OT_CKEY, UNITK=OT_UNITK, SM=OT_SM, RM=OT_RM, LAG=OT_LAG, OK=OT_OK, LATEK=OT_LATEK,
             UROW=OT_UROW, SROW=OT_SROW)
    t = dict(dname='单据日期', who='加工厂', want='外发加工厂', what='加工', pre='加工的')
    _put_rows(ws, ctx.get(SH_OUT, []), r0)
    for r in range(r0, r1 + 1):
        _bill_row(ws, r, c, hdr, '外发')
        ws[f'{OT_CHK}{r}'] = f'=IF({OT_SEQ}{r}="","",{_chain(_bill_checks(r, c, t), _s("√"))})'
    cols = [OT_SEQ, OT_DATE, OT_RDATE, OT_SUP, OT_NO, OT_OP, OT_STY, OT_ORD, OT_QTY, OT_UNIT, OT_PRICE, OT_AMTIN, OT_NOTE,
            OT_CMV, OT_LAGV, OT_CHK]
    autos = [OT_SEQ, OT_CMV, OT_LAGV, OT_CHK]
    style_rows(ws, r0, r1, cols, auto=autos,
               fmts={OT_DATE: DATE, OT_RDATE: DATE, OT_QTY: QTYF, OT_PRICE: MONEY, OT_AMTIN: MONEY, OT_CMV: '0"月"',
                     OT_LAGV: INTR, OT_NO: '@', OT_ORD: '@'},
               aligns={OT_SUP: AL, OT_NOTE: AL, OT_CHK: AL, OT_PRICE: AR, OT_AMTIN: AR, OT_QTY: AR},
               fills={col: FILL_IN for col in cols if col not in autos})
    heads = [(OT_SEQ, '序号'), (OT_DATE, '单据日期\n（回货/结算）'), (OT_RDATE, '收单日期\n（拿到单子）'), (OT_SUP, '加工厂'),
             (OT_NO, '单号'), (OT_OP, '加工内容'), (OT_STY, '用在哪\n（款式）'), (OT_ORD, '订单号\n（可空）'), (OT_QTY, '数量'),
             (OT_UNIT, '单位'), (OT_PRICE, '单价'), (OT_AMTIN, '金额\n（可空）'), (OT_NOTE, '备注'), (OT_CMV, '成本\n月份'),
             (OT_LAGV, '滞后\n天数'), (OT_CHK, '校验')]
    _heads(ws, hdr, heads, autos)
    _hidden(ws, hdr, r0, r1, [(OT_AMT, '金额'), (OT_CM, '成本月'), (OT_PDATE, '记账日期'), (OT_ATYPE, '款/公'), (OT_CKEY, '成本键'),
                               (OT_UNITK, '加工厂'), (OT_SM, '单据月'), (OT_RM, '收单月'), (OT_LAG, '滞后天数'), (OT_OK, '能记账'),
                               (OT_LATEK, '晚到键★'), (OT_UROW, '往来单位第几行'), (OT_SROW, '款式档案第几个')],
            fmts={OT_AMT: MONEY, OT_PDATE: DATE})
    ok, chk, amt_r = otr(OT_OK), otr(OT_CHK), otr(OT_AMT)
    _kpi(ws, 'A3:B3', '行数', 'C3', f'=COUNT({otr(OT_SEQ)})', INT)
    _kpi(ws, 'D3:E3', '本年加工费\n（能记账的）', 'F3:G3', f'=SUMIFS({amt_r},{ok},1)')
    _kpi(ws, 'H3:I3', '其中直接\n记到款式', 'J3:K3', f'=SUMIFS({amt_r},{ok},1,{otr(OT_ATYPE)},"款")')
    _kpi(ws, 'L3', '✗ 待改', 'M3', f'=COUNTIF({chk},"✗*")', INT)
    _kpi(ws, 'N3', '⚠ 提醒', 'O3', f'=COUNTIF({chk},"⚠*")', INT)
    put(ws, 'P3', f'="晚到（≥"&{LAGD}&"天）："&COUNT({otr(OT_LATEK)})&" 行"', F_KPI_L, fill('FFD9E1F2'), align=AL)
    ws.row_dimensions[3].height = 30
    _chk_cf(ws, OT_CHK, r0, r1)
    ws.conditional_formatting.add(f'{OT_LAGV}{r0}:{OT_LAGV}{r1}', FormulaRule(
        formula=[f'${OT_LATEK}{r0}<>""'], fill=fill('FFFCE4D6'), font=F_RED))
    dv_date(ws, f'{OT_DATE}{r0}:{OT_RDATE}{r1}')
    dv_list(ws, f'{OT_SUP}{r0}:{OT_SUP}{r1}', f'={UN_NAMES}', '【往来单位】登记过的外发加工厂')
    dv_list(ws, f'{OT_OP}{r0}:{OT_OP}{r1}', f'={OP_NAMES}', '【基础资料】⑨ 外发工序；没有的直接写', stop=False)
    dv_list(ws, f'{OT_STY}{r0}:{OT_STY}{r1}', f'={ST_CODES}', '这批加工是哪个款的；几个款一起没法分就空着', stop=False)
    ws.auto_filter.ref = f'A{hdr}:{OT_CHK}{r1}'
    ws.freeze_panes = f'E{r0}'


# ═══════════════════════════ 订单明细 ═══════════════════════════
# 隐藏列 AH SPK、AI NOK、AJ OFR、AK SMK、AL CQ 都在 layout 里登记了（OD_SPK … OD_CQ）


def build_od(wb, ctx):
    ws = wb[SH_ORD]
    hdr, r0, r1 = OD_HDR, OD_R0, OD_R1
    widths(ws, {'A': 5, 'B': 10, 'C': 11, 'D': 9, 'E': 10, 'F': 14, 'G': 7, 'H': 16, 'I': 11, 'J': 8, 'K': 9,
                'L': 9, 'M': 6, 'N': 11, 'O': 10, 'P': 11, 'Q': 11, 'R': 40, 'S': 10})
    title(ws, '订 单 明 细（电商部采购单 ＝ 客户订单 · 交货 · 每行的收入和成本）', OD_CHK, C_IN,
          '💡 电商部的采购单就是我们的订单。在电商部的表里只选数据行的「款式编码～备注」4 列（不要带标题行和最后一行合计），'
          '到这里 E 列空行右键 →「选择性粘贴 → 数值」（WPS：粘贴为数值），不要直接 Ctrl+V：采购单的备注是合并格、有黄底，直接粘会把合并格带进来，'
          '以后排序、复制行都会出错，款式编码的下拉也会被冲掉（粘过来的备注只有第一行有字）。每行补上订单号（采购单号）。'
          '交货了填「交货日期」（实交数量不填＝全交了）；收入按交货日期算到那个月。'
          '分批交货：复制这一行（看一眼只选了一行），第二行把订单数量清空，填第二次的交货日期和实交数量。'
          '电商部退回来的，另起一行实交填负数（只冲收入；这双的成本交货时已经算了，当返工 / 报废损失）；'
          '修好再交回去的再另起一行，填重交日期和双数，备注写「返修」（算收入，不再算交货双数和成本）。'
          '订单取消、不做了的，把订单数量改成实交数（不然这个款挂着的专用成本会一直按没交的双数留着，等后面交货）。'
          '某张单价钱特殊就填「结算单价」（不填按【款式档案】的结算单价）。客户空着＝【基础资料】的默认客户。'
          '款式编码要在【款式档案】登记过才有单价、才分得到成本（没登记的【款式档案】右边会列出来）。'
          '成本＝实交双数 × 这个款这个月的单双成本（自动算，晚到的送货单录进去以后会变）。不要插行删行（右边有隐藏公式）。')
    _put_rows(ws, ctx.get(SH_ORD, []), r0)
    sty_price = str_(ST_PRICE)
    spk_r, dq_r, oq_r, nok_r = odr(OD_SPK), odr(OD_DQ), odr(OD_OQ), odr(OD_NOK)
    digits = '{0,1,2,3,4,5,6,7,8,9}'
    for r in range(r0, r1 + 1):
        g = lambda col: f'{col}{r}'
        A, B, D, E, F, G, H, I, J, K = (g(x) for x in (OD_SEQ, OD_NO, OD_CUST, OD_STY, OD_SPEC, OD_QTY, OD_NOTE, OD_DDATE,
                                                       OD_DQTY, OD_PRICE))
        DQ, PU, DM, SROW, AMT, OK = (g(x) for x in (OD_DQ, OD_PRICEU, OD_DM, OD_SROW, OD_AMT, OD_OK))
        NOK, OFR, SMK, CQ = (g(x) for x in (OD_NOK, OD_OFR, OD_SMK, OD_CQ))
        i = f'ROW()-{hdr}'
        ws[A] = f'=IF(COUNTA({OD_NO}{r}:{OD_PRICE}{r})=0,"",{i})'
        ws[g(OD_CUSTK)] = f'=IF({A}="","",IF({_blank(D)},TRIM({CUST}&""),TRIM({D}&"")))'
        ws[DQ] = f'=IF({A}="",0,IF(ISNUMBER({I}),IF({_blank(J)},N({G}),N({J})),0))'
        ws[SROW] = _srow(E)
        ws[PU] = (f'=IF({A}="",0,IF(ISNUMBER({K}),{K},IF(NOT({_blank(K)}),0,'
                  f'IF({SROW}>0,N(INDEX({sty_price},{SROW})),0))))')
        ws[DM] = (f'=IF({A}="",0,IF(ISNUMBER({I}),IF({DQ}=0,0,IF(AND(YEAR({I})={YR},INT({I})>=N({OPEN})),MONTH({I}),0)),0))')
        ws[AMT] = f'=IF({DM}>0,ROUND({DQ}*{PU},2),0)'
        ws[g(OD_OQ)] = f'=IF({A}="",0,N({G}))'
        # 订单号规范写法（数字、文字、前后空格都算同一张单）；这张单第一次出现在第几行；订单首行★
        ws[NOK] = f'=TRIM({B}&"")'
        ws[OFR] = f'=IF({NOK}="",0,IFERROR(MATCH({esc(NOK)},{nok_r},0),0))'
        ws[g(OD_OFIRST)] = f'=IF({OFR}={i},{i},"")'
        osk = g(OD_OSKEY)
        ws[osk] = f'=IF(OR({NOK}="",{_blank(E)}),"",{NOK}&"|"&TRIM({E}&""))'
        ws[g(OD_OSFIRST)] = f'=IF({osk}="",0,IF(IFERROR(MATCH({esc(osk)},{odr(OD_OSKEY)},0),0)={i},1,0))'
        ws[g(OD_SNEW)] = (f'=IF(AND(NOT({_blank(E)}),{SROW}=0),'
                          f'IF(COUNTIF(${OD_STY}${hdr}:{OD_STY}{r - 1},TRIM({E}&""))=0,{i},""),"")')
        ws[g(OD_QSEL)] = f'=IF({NOK}="","",IF({NOK}=TRIM({OQ_SEL}&""),{i},""))'
        ws[OK] = f'=IF(AND({DM}>=1,{AMT}<>0,{g(OD_CHK)}<>"",LEFT({g(OD_CHK)},1)<>"✗"),1,0)'
        ws[g(OD_CSK)] = (f'=IF({OK}=1,IF(AND({g(OD_CUSTK)}=TRIM({CS_UNIT}&""),INT({I})>=N({CS_D1}),OR(N({CS_D2})=0,INT({I})<={CS_D2})),'
                         f'INT({I})*100000+{i},""),"")')
        # 款式×月数字键（_款式月 双数 / 收入按它 SUMIFS）；成本双数（退货、返修重交不算）
        ws[SMK] = f'=IF({OK}=1,{SROW}*16+{DM},0)'
        ws[CQ] = f'=IF({DQ}>0,IF(ISNUMBER(FIND("返修",{H}&"")),0,{DQ}),0)'
        spk = g(OD_SPK)
        ws[spk] = f'=IF({osk}="","",{osk}&"|"&SUBSTITUTE(TRIM({F}&""),"；",";"))'
        # 看得见的自动列
        s = f'SUBSTITUTE(TRIM({F}&""),"；",";")'
        ws[g(OD_COLOR)] = (f'=IF({_blank(F)},"",TRIM(LEFT({s},SUMPRODUCT(MIN(FIND({digits},{s}&"0123456789")))-1)))')
        ws[g(OD_SIZE)] = f'=IF(ISERROR(FIND(";",{s})),"",TRIM(MID({s},FIND(";",{s})+1,50)))'
        ws[g(OD_AMTV)] = f'=IF({A}="","",IF(ISNUMBER({I}),{AMT},""))'
        # 单双成本只给「成本双数」>0 的行（退货、返修重交没有）；SROW、DM 是 0 的不碰 INDEX
        ws[g(OD_UCV)] = (f'=IF({OK}<>1,"",IF({SROW}<1,"",IF({DM}<1,"",IF({CQ}<=0,"",'
                         f'N(INDEX({SM_UNIT2D},{SROW},{DM}))))))')
        ws[g(OD_COSTV)] = f'=IF({g(OD_AMTV)}="","",IF({g(OD_UCV)}="",0,ROUND({CQ}*{g(OD_UCV)},2)))'
        ws[g(OD_GPV)] = f'=IF({g(OD_AMTV)}="","",ROUND({g(OD_AMTV)}-{g(OD_COSTV)},2))'
        # 超交：只有补交行（订单数量空着）才按「订单号|款式|规格」整组 SUMIFS；普通行直接比本行 实交>订单数量
        sdq = f'SUMIFS({dq_r},{spk_r},{esc(spk)})'
        soq = f'SUMIFS({oq_r},{spk_r},{esc(spk)})'
        cust = g(OD_CUSTK)
        conds = [
            (_blank(B), _s('✗ 没填订单号（采购单号）')),
            (_blank(E), _s('✗ 没填款式编码')),
            (f'AND({_blank(G)},{_blank(J)})', _s('✗ 订单数量、实交数量都没填')),
            (_bad(I), _s('✗ 交货日期要填成日期（如 2026/9/12；没交就空着）')),
            (f'OR({_txtnum(G)},{_txtnum(J)},{_txtnum(K)})', _txtmsg('订单数量 / 实交数量 / 结算单价')),
            (f'OR({_bad(G)},{_bad(J)},{_bad(K)})', _s('✗ 订单数量、实交数量、结算单价要填数字')),
            (f'{SROW}=0', _s('⚠ 款式没在【款式档案】登记：没有结算单价、也分不到成本（【款式档案】右边列出来了，抄过去）')),
            (f'IF(ISNUMBER({I}),YEAR({I})<>{YR},FALSE)', _s('⚠ 交货日期不在本年度（不算本年收入）')),
            (f'IF(ISNUMBER({I}),INT({I})<N({OPEN}),FALSE)', _s('⚠ 建账以前交的货不算本账收入（欠的货款放【往来单位】期初应收）')),
            (f'AND(ISNUMBER({I}),{DQ}=0)', _s('⚠ 实交数量是 0（这行不算交货）')),
            (f'AND(ISNUMBER({I}),{AMT}=0)', _s('⚠ 没有结算单价（交货金额是 0）：在【款式档案】填结算单价，或在这行填')),
            (f'IF({DQ}>0,IF({_blank(G)},{sdq}>{soq},N({J})>N({G})),FALSE)',
             f'IF({_blank(G)},"⚠ 实交比订单数量多：这个订单这个款这个颜色码数一共交了 "&{sdq}&" 双，订了 "&{soq}&" 双",'
             f'"⚠ 实交比订单数量多：这行实交 "&N({J})&" 双，订了 "&N({G})&" 双")'),
            (f'IF({cust}="",TRUE,COUNTIF({UN_NAMES},{esc(cust)})=0)', _s('⚠ 客户不在【往来单位】（空着＝【基础资料】的默认客户）')),
            (f'NOT(ISNUMBER({I}))', _s('⏳ 还没交货')),
            (f'{DQ}<0', _s('√ 退货：只冲收入；成本交货时已算（当返工 / 报废损失）。修好再交回去的另起一行，备注写「返修」')),
            (f'AND({DQ}>0,{CQ}=0)', _s('√ 返修重交：算收入，不再算交货双数和成本')),
        ]
        ws[g(OD_CHK)] = f'=IF({A}="","",{_chain(conds, _s("√ 已交"))})'
    cols = [OD_SEQ, OD_NO, OD_DATE, OD_CUST, OD_STY, OD_SPEC, OD_QTY, OD_NOTE, OD_DDATE, OD_DQTY, OD_PRICE,
            OD_COLOR, OD_SIZE, OD_AMTV, OD_UCV, OD_COSTV, OD_GPV, OD_CHK]
    autos = [OD_SEQ, OD_COLOR, OD_SIZE, OD_AMTV, OD_UCV, OD_COSTV, OD_GPV, OD_CHK]
    style_rows(ws, r0, r1, cols, auto=autos,
               fmts={OD_NO: '@', OD_DATE: DATE, OD_DDATE: DATE, OD_QTY: INTR, OD_DQTY: INTR, OD_PRICE: MONEY,
                     OD_AMTV: MONEY, OD_UCV: MONEY, OD_COSTV: MONEY, OD_GPV: MONEY},
               aligns={OD_SPEC: AL, OD_NOTE: AL, OD_CHK: AL, OD_PRICE: AR, OD_AMTV: AR, OD_UCV: AR, OD_COSTV: AR, OD_GPV: AR},
               fills={col: FILL_IN for col in cols if col not in autos})
    heads = [(OD_SEQ, '序号'), (OD_NO, '订单号\n（采购单号）'), (OD_DATE, '下单日期'), (OD_CUST, '客户\n（空＝默认）'),
             (OD_STY, '款式编码'), (OD_SPEC, '颜色及规格\n（照采购单）'), (OD_QTY, '订单\n数量'), (OD_NOTE, '备注'),
             (OD_DDATE, '交货日期\n（空＝没交）'), (OD_DQTY, '实交数量\n（空＝全交）'), (OD_PRICE, '结算单价\n（空＝档案）'),
             (OD_COLOR, '颜色'), (OD_SIZE, '码数'), (OD_AMTV, '交货金额'), (OD_UCV, '单双成本'), (OD_COSTV, '成本'),
             (OD_GPV, '毛利'), (OD_CHK, '校验')]
    _heads(ws, hdr, heads, autos)
    _hidden(ws, hdr, r0, r1, [(OD_CUSTK, '客户'), (OD_DQ, '实交双数'), (OD_PRICEU, '用的单价'), (OD_DM, '交货月'),
                               (OD_SROW, '款式第几个'), (OD_AMT, '交货金额'), (OD_OQ, '订单数量'), (OD_OFIRST, '订单首行★'),
                               (OD_OSFIRST, '单内款式首行'), (OD_SNEW, '没登记款式★'), (OD_QSEL, '查询选中★'), (OD_OK, '能记账'),
                               (OD_CSK, '客户对账键'), (OD_OSKEY, '订单号|款式'), (OD_SPK, '订单|款式|规格'),
                               (OD_NOK, '订单号规范★'), (OD_OFR, '本单首行序号'), (OD_SMK, '款式×月键'), (OD_CQ, '成本双数★')],
            fmts={OD_AMT: MONEY, OD_PRICEU: MONEY, OD_NOK: '@'})
    ok, chk = odr(OD_OK), odr(OD_CHK)
    _kpi(ws, 'A3:B3', '订单数', 'C3', f'=COUNT({odr(OD_OFIRST)})', INT)
    _kpi(ws, 'D3', '订单双数', 'E3', f'=SUM({oq_r})', INT)
    _kpi(ws, 'F3', '已交双数', 'G3', f'=SUM({odr(OD_CQ)})', INT)          # 成本双数：不含退货、返修重交
    _kpi(ws, 'H3', '未交双数', 'I3', f'=MAX(0,E3-G3)', INT)
    _kpi(ws, 'J3:K3', '本年交货金额', 'L3:N3', f'=SUMIFS({odr(OD_AMT)},{ok},1)')
    _kpi(ws, 'O3', '✗ 待改', 'P3', f'=COUNTIF({chk},"✗*")', INT)
    _kpi(ws, 'Q3', '⚠ 提醒', 'R3', f'=COUNTIF({chk},"⚠*")', INT)
    ws.row_dimensions[3].height = 30
    _chk_cf(ws, OD_CHK, r0, r1)
    ws.conditional_formatting.add(f'{OD_GPV}{r0}:{OD_GPV}{r1}', FormulaRule(
        formula=[f'AND(ISNUMBER(${OD_GPV}{r0}),${OD_GPV}{r0}<0)'], font=F_RED))
    dv_date(ws, f'{OD_DATE}{r0}:{OD_DATE}{r1}')
    dv_date(ws, f'{OD_DDATE}{r0}:{OD_DDATE}{r1}')
    dv_list(ws, f'{OD_CUST}{r0}:{OD_CUST}{r1}', f'={UN_NAMES}', '空着＝【基础资料】的默认客户', stop=False)
    dv_list(ws, f'{OD_STY}{r0}:{OD_STY}{r1}', f'={ST_CODES}', '【款式档案】的款式编码；新款先照采购单写，再去档案登记', stop=False)
    ws.auto_filter.ref = f'A{hdr}:{OD_CHK}{r1}'
    ws.freeze_panes = f'C{r0}'


# ═══════════════════════════ 工资登记 ═══════════════════════════
WG_DROW, WG_SROW = 'Y', 'Z'             # 本表帮手列：部门在【基础资料】⑤ 第几行（0＝没有）；用在哪在【款式档案】第几个（0＝没登记）


def _by_type(x, k):
    """部门类型 → DP_TYPES 第 k 项（0 成本组件 1 借方科目 2 去向）"""
    f = '""'
    for tp, v in reversed(list(DP_TYPES.items())):
        f = f'IF({x}="{tp}","{v[k]}",{f})'
    return f


def build_wg(wb, ctx):
    ws = wb[SH_WAGE]
    hdr, r0, r1 = WG_HDR, WG_R0, WG_R1
    widths(ws, {'A': 5, 'B': 7, 'C': 10, 'D': 10, 'E': 11, 'F': 9, 'G': 9, 'H': 11, 'I': 11, 'J': 10, 'K': 10, 'L': 12,
                'M': 20, 'N': 10, 'O': 40, 'P': 10})
    title(ws, '工 资 登 记（每人每月一行 · 计件的可按款式分几行）', WG_CHK, C_IN,
          '💡 每人每月一行；计件的可以按款式分几行（同一个人同一个月几行没关系），「用在哪」填款式编码，这部分工资直接记到这个款'
          '（款式在【款式档案】改了「停产」的，它最后一次交货以后的月份按公共分）。'
          '「月份」填 1～12：工资算到那个月（不管哪天发的）。应发＝计件工资＋底薪/计时＋加班补贴－扣款。'
          '部门决定工资进哪里（【基础资料】⑤ 设）：直接生产（裁断、针车、成型、包装）＝直接人工，车间辅助（车间管理、仓管杂工）＝制造费用，'
          '这两种都分到款式成本里；办公室＝管理费用、业务跟单＝销售费用，不算成本。'
          '发工资在【资金日记账】记「付工资」。不要插行删行（右边有隐藏公式）。')
    _put_rows(ws, ctx.get(SH_WAGE, []), r0)
    for r in range(r0, r1 + 1):
        g = lambda col: f'{col}{r}'
        A, B, C, D, E = (g(x) for x in (WG_SEQ, WG_MON, WG_NAME, WG_DEPT, WG_STY))
        F_, G_, H_, I_, J_, K_ = (g(x) for x in (WG_PQ, WG_PP, WG_PIECE, WG_BASE, WG_OT, WG_DED))
        AMT, CM, DT, COMP, AT, OK, DROW, SROW = (g(x) for x in (WG_AMT, WG_CM, WG_DTYPE, WG_COMP, WG_ATYPE, WG_OK, WG_DROW,
                                                                WG_SROW))
        ws[A] = (f'=IF(COUNTA({WG_MON}{r}:{WG_PP}{r})+COUNTA({WG_BASE}{r}:{WG_DED}{r})+COUNTA({WG_NOTE}{r})=0,"",'
                 f'ROW()-{hdr})')
        ws[H_] = f'=IF({A}="","",IF(AND({_blank(F_)},{_blank(G_)}),"",ROUND(N({F_})*N({G_}),2)))'
        ws[AMT] = f'=IF({A}="",0,ROUND(N({H_})+N({I_})+N({J_})-N({K_}),2))'
        ws[g(WG_PAY)] = f'=IF({A}="","",{AMT})'
        ws[CM] = f'=IF({A}="",0,IF(ISNUMBER({B}),IF(AND({B}>=1,{B}<=12,{B}=INT({B})),IF({B}>={OM},{B},0),0),0))'
        ws[DROW] = f'=IF({_blank(D)},0,IFERROR(MATCH(TRIM({D}&""),{DP_NAMES},0),0))'
        ws[DT] = f'=IF({DROW}=0,"",INDEX({DP_TYPES_R},{DROW})&"")'
        ws[COMP] = f'={_by_type(DT, 0)}'
        ws[g(WG_DRC)] = f'={_by_type(DT, 1)}'
        ws[g(WG_DEST)] = f'=IF({A}="","",{_by_type(DT, 2)})'
        ws[SROW] = _srow(E)
        ws[AT] = f'=IF({A}="","",IF({COMP}="","公",{_atype(SROW, CM)}))'
        ws[g(WG_CKEY)] = f'=IF(AND({OK}=1,{COMP}<>""),IF({AT}="款","款|"&TRIM({E}&"")&"|"&{COMP},"公||"&{COMP}),"")'
        ws[OK] = f'=IF(AND({g(WG_CHK)}<>"",LEFT({g(WG_CHK)},1)<>"✗"),1,0)'
        conds = [
            (f'IF(ISNUMBER({B}),OR({B}<1,{B}>12,{B}<>INT({B})),TRUE)', _s('✗ 月份要填 1～12 的数字（工资属于几月）')),
            (f'{B}<{OM}', f'"✗ 月份早于建账月（"&{OM}&" 月）：建账前没发的工资放期初「应付职工薪酬」"'),
            (_blank(D), _s('✗ 没选部门')),
            (f'{DROW}=0', _s('✗ 部门不在【基础资料】⑤：先去加')),
            (f'{g(WG_DRC)}=""', _s('✗ 这个部门在【基础资料】⑤ 没选类型（直接生产/车间辅助/管理/销售）')),
            (f'OR({_txtnum(F_)},{_txtnum(G_)},{_txtnum(I_)},{_txtnum(J_)},{_txtnum(K_)})', _txtmsg('计件数量 / 单价 / 底薪 / 加班 / 扣款')),
            (f'OR({_bad(F_)},{_bad(G_)},{_bad(I_)},{_bad(J_)},{_bad(K_)})', _s('✗ 数字列填了文字：计件数量、单价、底薪、加班、扣款都要填数字')),
            (f'{AMT}=0', _s('✗ 应发是 0')),
            (_blank(C), _s('⚠ 没填姓名')),
            (f'{AMT}<0', _s('⚠ 应发是负数（扣的比发的多？）')),
            (f'AND(NOT({_blank(E)}),{COMP}="")', _s('⚠ 管理/销售部门填了用在哪（不起作用：不算款式成本）')),
            (f'AND(NOT({_blank(E)}),{AT}="公")',
             '"⚠ 用在哪的款式没在【款式档案】登记：先按公共分"'),
        ]
        ws[g(WG_CHK)] = f'=IF({A}="","",{_chain(conds, _s("√"))})'
    cols = [WG_SEQ, WG_MON, WG_NAME, WG_DEPT, WG_STY, WG_PQ, WG_PP, WG_PIECE, WG_BASE, WG_OT, WG_DED, WG_PAY, WG_NOTE,
            WG_DEST, WG_CHK]
    autos = [WG_SEQ, WG_PIECE, WG_PAY, WG_DEST, WG_CHK]
    style_rows(ws, r0, r1, cols, auto=autos,
               fmts={WG_MON: '0"月"', WG_PQ: QTYF, WG_PP: MONEY, WG_PIECE: MONEY, WG_BASE: MONEY, WG_OT: MONEY, WG_DED: MONEY,
                     WG_PAY: MONEY},
               aligns={WG_NOTE: AL, WG_CHK: AL, WG_PQ: AR, WG_PP: AR, WG_PIECE: AR, WG_BASE: AR, WG_OT: AR, WG_DED: AR, WG_PAY: AR},
               fills={col: FILL_IN for col in cols if col not in autos}, bold=[WG_PAY])
    heads = [(WG_SEQ, '序号'), (WG_MON, '月份\n（1～12）'), (WG_NAME, '姓名'), (WG_DEPT, '部门'), (WG_STY, '用在哪\n（款式，计件填）'),
             (WG_PQ, '计件数量'), (WG_PP, '计件单价'), (WG_PIECE, '计件工资'), (WG_BASE, '底薪 / 计时'), (WG_OT, '加班补贴'),
             (WG_DED, '扣款'), (WG_PAY, '应发工资'), (WG_NOTE, '备注'), (WG_DEST, '去向'), (WG_CHK, '校验')]
    _heads(ws, hdr, heads, autos)
    _hidden(ws, hdr, r0, r1, [(WG_AMT, '应发'), (WG_CM, '月份'), (WG_DTYPE, '部门类型'), (WG_COMP, '成本组件'), (WG_ATYPE, '款/公'),
                               (WG_CKEY, '成本键'), (WG_DRC, '借方科目'), (WG_OK, '能记账'), (WG_DROW, '部门第几行'),
                               (WG_SROW, '款式档案第几个')],
            fmts={WG_AMT: MONEY, WG_DRC: '@'})
    ok, chk, amt_r, drc = wgr(WG_OK), wgr(WG_CHK), wgr(WG_AMT), wgr(WG_DRC)
    _kpi(ws, 'A3:B3', '本年应发', 'C3:D3', f'=SUMIFS({amt_r},{ok},1)')
    for (lbl, val), (_tp, (_c, code, dest)) in zip((('E3', 'F3'), ('G3', 'H3'), ('I3', 'J3'), ('K3', 'L3')), DP_TYPES.items()):
        _kpi(ws, lbl, f'其中\n{dest}', val, f'=SUMIFS({amt_r},{ok},1,{drc},"{code}")')
    _kpi(ws, 'M3', '✗ 待改', 'N3', f'=COUNTIF({chk},"✗*")', INT)
    put(ws, 'O3', f'="⚠ 提醒："&COUNTIF({chk},"⚠*")&" 行"', F_KPI_L, fill('FFD9E1F2'), align=AL)
    ws.row_dimensions[3].height = 30
    _chk_cf(ws, WG_CHK, r0, r1)
    _dv_whole(ws, f'{WG_MON}{r0}:{WG_MON}{r1}', 1, 12, '填 1～12（工资属于几月）')
    dv_list(ws, f'{WG_DEPT}{r0}:{WG_DEPT}{r1}', f'={DP_NAMES}', '【基础资料】⑤ 的部门')
    dv_list(ws, f'{WG_STY}{r0}:{WG_STY}{r1}', f'={ST_CODES}', '计件按款式记才填；计时的空着', stop=False)
    ws.auto_filter.ref = f'A{hdr}:{WG_CHK}{r1}'
    ws.freeze_panes = f'D{r0}'


def build(wb, ctx):
    build_dn(wb, ctx)
    build_ot(wb, ctx)
    build_od(wb, ctx)
    build_wg(wb, ctx)
