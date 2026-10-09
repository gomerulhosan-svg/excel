# -*- coding: utf-8 -*-
"""录入表（蓝）：收支登记（登记汇总表）、应收应付登记；基础资料（灰）：基础资料、往来单位。
   录入表只放录入列，外加最右边两列只给人看的显示列（别的公式不引用它们，插进来的行没有这两格也不影响汇总）。"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *


def _d(v):
    if v in (None, ''):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.datetime(v.year, v.month, v.day)
    try:
        return dt.datetime.strptime(str(v)[:10], '%Y-%m-%d')
    except ValueError:
        return v


def dv_dec(ws, sqref, what='金额'):
    dv = DataValidation(type='decimal', operator='between', formula1='-999999999', formula2='999999999', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what, error=f'{what}要填数字（不要带文字、单位）')
    ws.add_data_validation(dv)
    dv.add(sqref)


QTY = '#,##0;[Red]-#,##0;"-"'
PRICE = '#,##0.00##'
FILL_EX = fill('FFEDEDED')


def _example_grey(ws, rows, r0, last_col):
    """示例行（备注以「示例」开头）灰底，正式用时整行删掉"""
    for i, x in enumerate(rows):
        if str(x.get('备注') or '').startswith('示例'):
            for c in range(1, CI(last_col) - 1):
                ws.cell(row=r0 + i, column=c).fill = FILL_EX


def dv_len(ws, sqref, n=50):
    dv = DataValidation(type='textLength', operator='lessThanOrEqual', formula1=str(n), allow_blank=True, showErrorMessage=True,
                        errorStyle='stop', errorTitle='名称太长', error=f'名称不要超过 {n} 个字')
    ws.add_data_validation(dv)
    dv.add(sqref)


def _frame(ws, last_col, hdr, heads, color, tip, W, r0, nstyle, fm, al, freeze):
    widths(ws, W)
    title(ws, ws.title, last_col, color, tip)
    header(ws, hdr, heads, color)
    style_rows(ws, r0, r0 + nstyle - 1, [c for c, _ in heads], fmts=fm, aligns=al)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f'A{hdr}:{last_col}{r0 + nstyle - 1}'


def _write_rows(ws, rows, cols, r0, date_keys=('日期',), num_keys=()):
    for i, x in enumerate(rows):
        r = r0 + i
        for k, col in cols.items():
            v = x.get(k)
            if v in (None, ''):
                continue
            if k in date_keys:
                v = _d(v)
            elif v == 'QP':                      # 金额＝数量×单价（跟用户原表一样写公式）
                v = f'=ROUND({cols["数量"]}{r}*{cols["单价"]}{r},2)'
            ws[f'{col}{r}'] = v


# ───────────────────────── 收支登记（登记汇总表） ─────────────────────────
def build_cash(ws, ctx):
    C = CASH_COLS
    S = CASH_SHOW
    heads = [(C['日期'], '日期'), (C['板块'], '业务板块'), (C['收支项目'], '收支项目'), (C['摘要'], '摘要'),
             (C['往来单位'], '往来单位\n（客户/供应商）'), (C['数量'], '数量\n（KG）'), (C['单价'], '单价'),
             (C['收入'], '收入金额'), (C['支出'], '支出金额'), (C['账户'], '账户\n（空＝现金）'), (C['经手人'], '经手人'),
             (C['备注'], '备注'), (S['结余'], '本账户结余\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 登记汇总表：所有板块的钱都在这一张表登记，一笔一行：日期、业务板块、收支项目都从下拉选，各板块表、查询、资金台账自动拆分汇总。灰底的是示例行，正式用时整行删掉。'
           '收入/支出填金额；粮食买卖、烘干按斤算的，也可以只填数量和单价，金额空着就按「数量×单价」算（收支项目是收入类算收入，支出类算支出）。'
           '收回以前的欠款选「收回欠款」、付掉欠人家的钱选「支付欠款」，往来单位一定要选，应收应付和对账单会自动冲掉。'
           '现金存银行、银行取现：记两行「内部转账」（一行支出、一行收入），不算收支。可以插行、删行、排序（整行一起动）。')
    W = {C['日期']: 11, C['板块']: 12, C['收支项目']: 12, C['摘要']: 24, C['往来单位']: 13, C['数量']: 10, C['单价']: 9,
         C['收入']: 13, C['支出']: 13, C['账户']: 10, C['经手人']: 8, C['备注']: 14, S['结余']: 14, S['校验']: 26}
    fm = {C['日期']: DATE, C['数量']: QTY, C['单价']: PRICE, C['收入']: MONEY, C['支出']: MONEY, S['结余']: MONEY}
    al = {C['摘要']: AL, C['备注']: AL, C['收入']: AR, C['支出']: AR, S['结余']: AR, S['校验']: AL, C['数量']: AR, C['单价']: AR}
    rows = ctx.get('cash', [])
    nshow = min(max(len(rows) + 1000, 1500), N_CASH)
    _frame(ws, CASH_LAST, CASH_HDR, heads, C_IN, tip, W, CASH_R0, nshow, fm, al, f'D{CASH_R0}')
    ws['A1'].value = '收支登记（登记汇总表）'
    ws.row_dimensions[CASH_HDR].height = 40
    # 第 3 行：合计。左边＝筛选后看得见的行里填了金额的（SUBTOTAL，区域从表头行起，插在第一行上面也算进去）；右边＝全部（含按数量×单价算的）
    last = CASH_R0 + N_CASH + 2000
    put(ws, 'A3', '筛选后看得见的行（填了金额的）', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('A3:G3')
    for col in (C['收入'], C['支出']):
        put(ws, f'{col}3', f'=SUBTOTAL(9,{col}{CASH_HDR}:{col}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["账户"]}3', '全部收入', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, f'{C["经手人"]}3', '=SUM(收_收入)', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["备注"]}3', '全部支出', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, f'{S["结余"]}3', '=SUM(收_支出)', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{S["校验"]}3', '←「全部」含金额空着按数量×单价算的', F_NOTE, align=AL, border=False)
    # 数据
    _write_rows(ws, rows, C, CASH_R0)
    # 显示列
    base = f'ROW()-ROW(${C["日期"]}${CASH_HDR})'
    for i in range(nshow):
        r = CASH_R0 + i
        n = base
        ws[f'{S["结余"]}{r}'] = (f'=IFERROR(IF(INDEX(收_有效,{n})<>1,"",ROUND(SUMIFS(账户_期初余额,账户_名称,INDEX(收_账户,{n}))'
                                f'+SUMIFS(收_净额,收_账户,INDEX(收_账户,{n}),收_排序键,"<="&INDEX(收_排序键,{n})),2)),"")')
        ws[f'{S["校验"]}{r}'] = f'=IFERROR(INDEX(收_校验,{n})&"","")'
        for col in (S['结余'], S['校验']):
            ws[f'{col}{r}'].fill = FILL_AUTO
            ws[f'{col}{r}'].font = F_AUTO
    ws.conditional_formatting.add(f'{S["校验"]}{CASH_R0}:{S["校验"]}{CASH_R0 + nshow - 1}',
                                  FormulaRule(formula=[f'LEFT({S["校验"]}{CASH_R0},1)="✗"'], font=F_RED))
    # 下拉、格式校验（覆盖到容量）
    rng = lambda col: f'{col}{CASH_R0}:{col}{CASH_R0 + N_CASH - 1}'
    dv_date(ws, rng(C['日期']))
    dv_list(ws, rng(C['板块']), '=板块列表', prompt='从【基础资料】的业务板块里选；账户之间倒钱可以不选')
    dv_list(ws, rng(C['收支项目']), '=收支项目列表', prompt='收回以前的欠款选「收回欠款」，付掉欠人家的钱选「支付欠款」')
    dv_list(ws, rng(C['往来单位']), '=往来单位列表', stop=False, prompt='客户、供应商从【往来单位】选；收回/支付欠款一定要选')
    dv_list(ws, rng(C['账户']), '=账户列表', prompt='空着＝基础资料里第一个账户（现金）')
    dv_list(ws, rng(C['经手人']), '=经手人列表', stop=False)
    for k in ('数量', '单价'):
        dv_dec(ws, rng(C[k]), k)
    for k in ('收入', '支出'):
        dv_dec(ws, rng(C[k]))
    _example_grey(ws, rows, CASH_R0, CASH_LAST)
    home_link(ws, f'{CL(CI(CASH_LAST) + 1)}1')
    print_setup(ws, f'{CASH_HDR}:{CASH_HDR}')


# ───────────────────────── 应收应付登记 ─────────────────────────
def build_wl(ws, ctx):
    C = WL_COLS
    S = WL_SHOW
    heads = [(C['日期'], '日期'), (C['板块'], '业务板块'), (C['往来单位'], '往来单位'), (C['类型'], '应收 / 应付'),
             (C['收支项目'], '业务内容\n（收支项目）'), (C['摘要'], '摘要'), (C['数量'], '数量\n（KG）'), (C['单价'], '单价'),
             (C['金额'], '金额\n（空＝数量×单价）'), (C['约定日期'], '约定收/付款日期\n（空＝按账期）'), (C['经手人'], '经手人'),
             (C['备注'], '备注'), (S['未结'], '这笔还没结的\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 赊出去、赊进来（钱当时没收/没付）的业务在这里登记：卖给人家没收钱＝应收，买人家的没给钱＝应付。灰底的是示例行，正式用时整行删掉。'
           '金额空着就按「数量×单价」算；折让、少收少付记负数。以后收到钱/付了钱，到【收支登记】记一笔「收回欠款」「支付欠款」并选同一个往来单位，'
           '这里「这笔还没结的」会按先欠先还自动减掉。建账以前的老欠款也在这里记一行（摘要写「期初余额」）。可以插行、删行。')
    W = {C['日期']: 11, C['板块']: 12, C['往来单位']: 14, C['类型']: 8, C['收支项目']: 12, C['摘要']: 24, C['数量']: 10,
         C['单价']: 9, C['金额']: 13, C['约定日期']: 13, C['经手人']: 8, C['备注']: 16, S['未结']: 13, S['校验']: 32}
    fm = {C['日期']: DATE, C['约定日期']: DATE, C['数量']: QTY, C['单价']: PRICE, C['金额']: MONEY, S['未结']: MONEY}
    al = {C['摘要']: AL, C['备注']: AL, C['金额']: AR, S['未结']: AR, S['校验']: AL, C['数量']: AR, C['单价']: AR}
    rows = ctx.get('wl', [])
    nshow = min(max(len(rows) + 500, 800), N_WL)
    _frame(ws, WL_LAST, WL_HDR, heads, C_IN, tip, W, WL_R0, nshow, fm, al, f'D{WL_R0}')
    ws.row_dimensions[WL_HDR].height = 40
    for c, lbl, f in (('A', '应收合计（全部）', '=SUM(往_应收额)'), ('D', '应付合计（全部）', '=SUM(往_应付额)'),
                      ('G', '应收还没结', '=SUMIFS(往_未结,往_类型,"应收")'), ('J', '应付还没结', '=SUMIFS(往_未结,往_类型,"应付")')):
        put(ws, f'{c}3', lbl, F_KPI_L, fill('FFD9E1F2'), align=AC)
        ws.merge_cells(f'{c}3:{CL(CI(c) + 1)}3')
        put(ws, f'{CL(CI(c) + 2)}3', f, F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, 'M3', '还没结的按首页截止日算', F_NOTE, align=AL, border=False)
    _write_rows(ws, rows, C, WL_R0, date_keys=('日期', '约定日期'))
    base = f'ROW()-ROW(${C["日期"]}${WL_HDR})'
    for i in range(nshow):
        r = WL_R0 + i
        ws[f'{S["未结"]}{r}'] = f'=IFERROR(IF(INDEX(往_有效,{base})<>1,"",INDEX(往_未结,{base})),"")'
        ws[f'{S["校验"]}{r}'] = f'=IFERROR(INDEX(往_校验,{base})&"","")'
        for col in (S['未结'], S['校验']):
            ws[f'{col}{r}'].fill = FILL_AUTO
            ws[f'{col}{r}'].font = F_AUTO
    ws.conditional_formatting.add(f'{S["校验"]}{WL_R0}:{S["校验"]}{WL_R0 + nshow - 1}',
                                  FormulaRule(formula=[f'LEFT({S["校验"]}{WL_R0},1)="✗"'], font=F_RED))
    rng = lambda col: f'{col}{WL_R0}:{col}{WL_R0 + N_WL - 1}'
    dv_date(ws, rng(C['日期']))
    dv_date(ws, rng(C['约定日期']))
    dv_list(ws, rng(C['板块']), '=板块列表')
    dv_list(ws, rng(C['往来单位']), '=往来单位列表', prompt='从【往来单位】选；新单位先去那里加一行')
    dv_list(ws, rng(C['类型']), '"' + ','.join(WL_TYPES) + '"', prompt='卖出去没收钱＝应收；买进来没付钱＝应付')
    dv_list(ws, rng(C['收支项目']), '=收支项目列表', stop=False)
    dv_list(ws, rng(C['经手人']), '=经手人列表', stop=False)
    for k in ('数量', '单价'):
        dv_dec(ws, rng(C[k]), k)
    dv_dec(ws, rng(C['金额']))
    _example_grey(ws, rows, WL_R0, WL_LAST)
    home_link(ws, f'{CL(CI(WL_LAST) + 1)}1')
    print_setup(ws, f'{WL_HDR}:{WL_HDR}')


# ───────────────────────── 基础资料 ─────────────────────────
def build_base(ws, ctx):
    widths(ws, {'A': 16, 'B': 13, 'C': 14, 'D': 13, 'E': 30, 'F': 3, 'G': 60})
    title(ws, '基 础 资 料', 'G', C_MASTER,
          '💡 一次填好、偶尔改：① 单位名称、建账日期；② 业务板块（各板块表、查询按它分）；③ 账户（现金、银行，填建账日的余额）；'
          '④ 收支项目（收支登记的下拉，「方向」「用途」「库存」决定怎么算）；⑤ 经手人。往来单位在【往来单位】那张表。'
          '每块上下排着，在块里插整行、删整行都只动这一块；名称不要重复。改了板块名、项目名，记得用「查找替换」把两张登记表里的旧名字也改掉。')
    home_link(ws, 'G3')
    pv = ctx.get('params', {})
    put(ws, 'A3', '① 参数', F_SEC, fill(C_MASTER), align=AL)
    ws.merge_cells('A3:C3')
    for k, lbl, default in PARAMS:
        r = PAR_ROW[k]
        put(ws, f'A{r}', lbl, F_KPI_L, fill('FFD9E1F2'), align=AL)
        ws.merge_cells(f'A{r}:B{r}')
        v = pv.get(k) or default
        if k == '建账日':
            v = _d(v)
        put(ws, f'C{r}', v, F_IN, FILL_IN, DATE if k == '建账日' else None, AC)
    dv_date(ws, f'C{PAR_ROW["建账日"]}')
    put(ws, f'E{PAR_ROW["单位名称"]}', '对账单抬头用，请填单位全称', F_NOTE, align=AL, border=False)
    put(ws, f'E{PAR_ROW["建账日"]}', '从哪天开始用这本账；期初余额都是这一天的', F_NOTE, align=AL, border=False)
    put(ws, f'E{PAR_ROW["账期"]}', '赊账没填约定日期时，登记日期＋这么多天算到期', F_NOTE, align=AL, border=False)
    notes = {'板块': '期初结余＝建账日这个板块手上的钱；粮食购销填期初粮食存量（KG）；有设备房屋的填期初固定资产原值。',
             '账户': '空着的账户栏＝现金。期初余额填建账日的余额。',
             '项目': '方向：收入/支出/双向。用途：普通；冲应收＝收回欠款；冲应付＝支付欠款；固定资产＝买设备房屋；内部转账＝账户间倒钱（不算收支）；'
                     '借款、调拨＝照常算收支，首页单独列出来（不是经营收支）。库存：入库＝粮食购进，出库＝粮食销售（按 KG）。',
             '经手人': '收支登记、应收应付登记的经手人从这里下拉；查询、板块表可以按经手人看。'}
    heads = {'板块': (SEG_COLS, [('名称', '板块名称'), ('期初结余', '期初结余\n（建账日）'), ('期初存量', '期初粮食存量\n（KG）'),
                                 ('期初固定资产', '期初固定资产\n（原值）'), ('备注', '备注')], ctx.get('segments', [])),
             '账户': (ACC_COLS, [('名称', '账户名称'), ('期初余额', '期初余额\n（建账日）'), ('备注', '备注')], ctx.get('accounts', [])),
             '项目': (ITEM_COLS, [('名称', '收支项目'), ('方向', '方向'), ('用途', '用途'), ('库存', '库存'), ('备注', '说明')],
                      ctx.get('items', [])),
             '经手人': (PER_COLS, [('姓名', '经手人'), ('备注', '备注')], ctx.get('persons', []))}
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        cols, hd, data = heads[key]
        c2 = cols[hd[-1][0]]
        put(ws, f'A{tr}', ttl, F_SEC, fill(C_MASTER), align=AL)
        ws.merge_cells(f'A{tr}:{c2}{tr}')
        put(ws, f'G{tr}', notes[key], F_NOTE, align=ALW, border=False)
        ws.merge_cells(f'G{tr}:G{hr}')
        header(ws, hr, [(cols[k], t) for k, t in hd], C_MASTER, height=34)
        fm = {cols[k]: MONEY for k in ('期初结余', '期初固定资产', '期初余额') if k in cols}
        if '期初存量' in cols:
            fm[cols['期初存量']] = QTY
        r0 = hr + 1
        style_rows(ws, r0, r0 + cap - 1, [cols[k] for k, _ in hd], fmts=fm,
                   aligns={cols[k]: AL for k in ('名称', '备注', '姓名') if k in cols})
        for i, x in enumerate(data):
            for k, _t in hd:
                v = x.get(k)
                if v not in (None, ''):
                    ws[f'{cols[k]}{r0 + i}'] = v
        for k in ('期初结余', '期初存量', '期初固定资产', '期初余额'):
            if k in cols:
                dv_dec(ws, f'{cols[k]}{r0}:{cols[k]}{r0 + cap - 1}')
        dv_len(ws, f'A{r0}:A{r0 + cap - 1}', 30)
        if key == '项目':
            rngI = lambda k: f'{ITEM_COLS[k]}{r0}:{ITEM_COLS[k]}{r0 + cap - 1}'
            dv_list(ws, rngI('方向'), '"' + ','.join(ITEM_DIRS) + '"')
            dv_list(ws, rngI('用途'), '"' + ','.join(ITEM_USES) + '"', prompt='空着＝普通')
            dv_list(ws, rngI('库存'), '"' + ','.join(ITEM_STOCK) + '"', prompt='只有粮食购进（入库）、粮食销售（出库）要选')
        put(ws, f'E{r0 + cap}', f'（这一块最多 {cap} 个，超过了【数据校验】会提醒）', F_NOTE, align=AL, border=False)
    ws.freeze_panes = 'A4'
    print_setup(ws, None, landscape=False)


# ───────────────────────── 往来单位 ─────────────────────────
def build_unit(ws, ctx):
    C = UNIT_COLS
    heads = [(C['名称'], '单位 / 个人名称'), (C['类型'], '类型'), (C['板块'], '主要业务板块'), (C['联系人'], '联系人'),
             (C['电话'], '电话'), (C['地址'], '地址'), (C['账期'], '账期（天）\n（空＝默认）'), (C['备注'], '备注')]
    tip = ('💡 客户、供应商（单位或个人）一个一行，名称不要重复。收支登记、应收应付登记的「往来单位」从这里下拉；'
           '应收应付跟进、对账单按这个名称汇总。账期＝赊账多少天内要结清，空着按【基础资料】的默认账期。备注可以写跟进情况（应收应付跟进表会显示）。可以插行、删行。')
    W = {C['名称']: 18, C['类型']: 12, C['板块']: 13, C['联系人']: 10, C['电话']: 14, C['地址']: 26, C['账期']: 11, C['备注']: 24}
    rows = ctx.get('units', [])
    _frame(ws, 'H', UNIT_HDR, heads, C_MASTER, tip, W, UNIT_R0, max(len(rows) + 50, 100), {},
           {C['名称']: AL, C['地址']: AL, C['备注']: AL}, f'B{UNIT_R0}')
    _write_rows(ws, rows, C, UNIT_R0, date_keys=())
    rng = lambda col: f'{col}{UNIT_R0}:{col}{UNIT_R0 + N_UNIT - 1}'
    dv_list(ws, rng(C['类型']), '"' + ','.join(UNIT_TYPES) + '"')
    dv_list(ws, rng(C['板块']), '=板块列表', stop=False)
    dv = DataValidation(type='whole', operator='between', formula1='0', formula2='3650', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng(C['账期']))
    dv_len(ws, rng(C['名称']), 50)
    home_link(ws, 'I1')
    print_setup(ws, f'{UNIT_HDR}:{UNIT_HDR}')


BUILDERS = {SH_CASH: build_cash, SH_WL: build_wl, SH_BASE: build_base, SH_UNIT: build_unit}


def build(wb, ctx):
    for name, fn in BUILDERS.items():
        fn(wb[name], ctx)
