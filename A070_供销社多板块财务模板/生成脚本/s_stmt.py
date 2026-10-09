# -*- coding: utf-8 -*-
"""【对账单】（往来组，可直接打印给对方签字）：黄格选往来单位（下拉，也可以手打名单外的）、板块（空＝全部）、起（空＝建账日）、止（空＝首页截止日）。
   抬头：单位名称＋「对 账 单」、往来单位、联系人、电话、对账期间、板块。
   汇总：期初（起以前）应收/应付余额、本期增加、本期收款/付款、期末余额；净额（对方欠我们为正）。
   明细（按排序键：日期，同一天收支登记在前，最多 500 笔）：第一行期初余额，逐笔滚动应收余额、应付余额；合计行＋核对；签章。
   口径：设计.md 第 3 节「对账单」——应收应付登记的行：应收→应收增加、应付→应付增加；
   收支登记里往来单位＝它的行：冲应收→收款、冲应付→付款；其他（现结）：收入→应收增加＝收款，支出→应付增加＝付款；内部转账不列。
   明细是「活动区」：隐藏列先算出每一行是什么（期初 / 明细 / 合计 / 共几笔 / 签章），合计和签章紧跟在最后一笔后面，打印区域也跟着伸缩。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, Border, Side, PatternFill, Alignment
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'L'
NL = 500
R_HEAD, R_TO, R_CT, R_NOTE = 6, 7, 8, 9
R_SH, R_SAR, R_SAP, R_SNET = 10, 11, 12, 13
R_HDR, R0 = 15, 16
NREG = NL + 12
R1 = R0 + NREG - 1
# 隐藏列
SL, SV = 'O', 'P'
SC = ['单位', '单位条件', '板块', '板块条件', '起', '止', '名单序号', '笔数', 'n', '最后一行', '核对', '现结收', '现结付']
SR = {k: 3 + i for i, k in enumerate(SC)}
HK, HC, HKEY, HN, HSRC, HUSE = 'R', 'S', 'T', 'U', 'V', 'W'
# 排序键折成 500 行一列的方块（SMALL 可以对整块取第 k 小），隐藏列都在活动区的行数以内，打印、滚动都不会拖出几千行
BLK = 500
assert N_WL % BLK == 0 and N_CASH % BLK == 0 and BLK <= NREG
YN = 'Y'                                                        # 方块的行号 1～500（第 n 条＝行号＋列偏移）
YKS = [CL(CI('Z') + j) for j in range(N_CASH // BLK)]           # _收 的排序键（Z..AI）
YKW = [CL(CI(YKS[-1]) + 1 + j) for j in range(N_WL // BLK)]     # _往 的排序键（AJ..AM）
KEYS = f'${YKS[0]}${R0}:${YKW[-1]}${R0 + BLK - 1}'

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_BIG = Font(name=YH, sz=16, bold=True, color='FF000000')
F_HEADB = Font(name=YH, sz=11, bold=True, color='FF000000')
F_SIGN = Font(name=YH, sz=11, bold=True, color='FF000000')
F_GRAY = Font(name=YH, sz=9, color='FF808080')
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_SUM_B = Font(name=YH, sz=10, bold=True, color='FF000000')
F_SUM_V = Font(name=YH, sz=11, bold=True, color='FF1F3864')


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def _sw(cd, m):
    return '=' + _chain([(f'{cd}="{k}"', v) for k, v in m.items()])


TIP = ('💡 给客户、供应商对账用，可以直接打印（A4 横向）请对方签字盖章。黄格：往来单位（从下拉选；名单里没有的也可以直接打名字）、'
       '板块（空＝全部板块）、起、止（空着＝建账日～首页截止日，右边灰格是实际用的日期）。'
       '上面是汇总：期初余额（起日以前）＋本期增加－本期收款/付款＝期末余额；下面逐笔列出（同一天先列收支登记的），应收、应付余额逐笔往下滚，合计跟汇总自动核对。'
       '赊账在【应收应付登记】记，收回/支付欠款和现结（当场结清的买卖，增加和收付款同时记，余额不变）在【收支登记】记，要改去那边改。')


def build(wb, ctx=None):
    ctx = ctx or {}
    ws = wb[SH_STMT]
    W = {'A': 12.5, 'B': 12.5, 'C': 28, 'D': 11, 'E': 9, 'F': 13.5, 'G': 13.5, 'H': 15, 'I': 13.5, 'J': 13.5, 'K': 15,
         'L': 18, 'M': 10}
    widths(ws, W)
    title(ws, '对账单', LAST, C_WL, TIP)
    home_link(ws, 'M1')

    # ── 第 3、4 行：选择格 ──
    units = [u.get('名称') for u in ctx.get('units', []) if u.get('名称')]
    selector(ws, 'A3', '往来单位', 'B3', units[0] if units else None, '=往来单位列表',
             prompt='从下拉选；【往来单位】名单里没有的也可以直接打名字')
    ws.merge_cells('B3:C3')
    selector(ws, 'D3', '板块', 'E3', None, '=板块列表', prompt='空＝全部板块')
    ws.merge_cells('E3:F3')
    selector(ws, 'G3', '起', 'H3', None, fmt=DATE, prompt='空＝建账日')
    selector(ws, 'I3', '止', 'J3', None, fmt=DATE, prompt='空＝首页截止日')
    dv_date(ws, 'H3')
    dv_date(ws, 'J3')
    ws.merge_cells('K3:L3')
    put(ws, 'K3', '起止空着＝建账日～首页截止日', F_NOTE, align=AL, border=False)
    put(ws, 'A4', '实际用的', F_NOTE, align=AC, border=False)
    ws.merge_cells('B4:C4')
    put(ws, 'B4', (f'=IF({S("单位")}="","← 先从下拉选往来单位",IF({S("名单序号")}>0,"✓ "&IF(INDEX(单位_类型,{S("名单序号")})="",'
                   f'"在往来单位名单里",INDEX(单位_类型,{S("名单序号")})),"⚠ 不在【往来单位】名单里（照样能出）"))'),
        F_AUTOB, FILL_AUTO, align=AC)
    ws.merge_cells('E4:F4')
    put(ws, 'E4', f'=IF({S("板块")}="","全部板块",IF(ISNA(MATCH({S("板块")},板块_名称,0)),"⚠ 不在基础资料里",{S("板块")}))',
        F_AUTOB, FILL_AUTO, align=AC)
    put(ws, 'H4', f'={S("起")}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'J4', f'={S("止")}', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.merge_cells('K4:L4')
    put(ws, 'K4', (f'=IF({S("起")}>{S("止")},"✗ 起 晚于 止，请改日期",IF({S("笔数")}>{NL},"⚠ 超过 {NL} 笔，只列了前 {NL} 笔，请缩短日期",'
                   f'IF(AND({S("单位")}<>"",{S("核对")}=0),"✗ 明细跟汇总对不上","")))'), F_WARN, align=AL, border=False)
    ws.row_dimensions[3].height = 22
    ws.row_dimensions[4].height = 20
    ws.row_dimensions[5].height = 8

    # ── 标量 ──
    U, EU, SG, S0, S1 = S('单位'), S('单位条件'), S('板块条件'), S('起'), S('止')
    per = lambda nm: f'{nm},">="&{S0},{nm},"<="&{S1}'
    bef = lambda nm: f'{nm},"<"&{S0}'
    xj = lambda col, cond: (f'SUMIFS({col},收_往来单位,{EU},收_板块,{SG},收_内部转账,0,收_用途,"<>冲应收",收_用途,"<>冲应付",{cond})')
    rk = lambda c: f'${c}${R0}:${c}${R1}'
    tot = lambda c: f'SUMIFS({rk(c)},{rk(HC)},"T")'
    sc = {
        '单位': '=TRIM(B3&"")',
        '单位条件': f'={esc(U)}',
        '板块': '=TRIM(E3&"")',
        '板块条件': f'=IF({S("板块")}="","<>@@全部@@",{esc(S("板块"))})',
        '起': '=IF(ISNUMBER(H3),INT(H3),P_建账日)',
        '止': '=IF(ISNUMBER(J3),INT(J3),P_截止)',
        '名单序号': f'=IF({U}="",0,IFERROR(MATCH({EU},单位_名称,0),0))',
        '笔数': f'=COUNT({KEYS})',
        'n': f'=MIN({S("笔数")},{NL})',
        '最后一行': f'={R0 - 1}+{S("n")}+12',
        '核对': (f'=IF(AND(ROUND({tot("F")}-G{R_SAR},2)=0,ROUND({tot("G")}-H{R_SAR},2)=0,ROUND({tot("H")}-I{R_SAR},2)=0,'
                 f'ROUND({tot("I")}-G{R_SAP},2)=0,ROUND({tot("J")}-H{R_SAP},2)=0,ROUND({tot("K")}-I{R_SAP},2)=0),1,0)'),
        '现结收': f'=IF({U}="",0,ROUND({xj("收_收入", per("收_日期"))},2))',
        '现结付': f'=IF({U}="",0,ROUND({xj("收_支出", per("收_日期"))},2))',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    # ── 抬头 ──
    ws.merge_cells(f'A{R_HEAD}:{LAST}{R_HEAD}')
    put(ws, f'A{R_HEAD}', '=IF(P_单位名称="","",P_单位名称&"　")&"对 账 单"', F_BIG, align=AC, border=False)
    ws.row_dimensions[R_HEAD].height = 32
    ws.merge_cells(f'A{R_TO}:F{R_TO}')
    put(ws, f'A{R_TO}', f'="往来单位："&IF({U}="","（还没选）",{U})', F_HEADB, align=AL, border=False)
    ws.merge_cells(f'G{R_TO}:{LAST}{R_TO}')
    put(ws, f'G{R_TO}', f'="对账期间："&TEXT({S0},"yyyy-mm-dd")&" 至 "&TEXT({S1},"yyyy-mm-dd")', F_HEADB, align=AR, border=False)
    ws.merge_cells(f'A{R_CT}:F{R_CT}')
    ix = S('名单序号')
    put(ws, f'A{R_CT}', (f'="联系人："&IF({ix}=0,"",INDEX(单位_联系人,{ix}))&"　　　电话："&IF({ix}=0,"",INDEX(单位_电话,{ix}))'),
        F_TXT, align=AL, border=False)
    ws.merge_cells(f'G{R_CT}:{LAST}{R_CT}')
    put(ws, f'G{R_CT}', f'="业务板块："&IF({S("板块")}="","全部",{S("板块")})', F_TXT, align=AR, border=False)
    ws.merge_cells(f'A{R_NOTE}:{LAST}{R_NOTE}')
    put(ws, f'A{R_NOTE}', '金额单位：元。应收＝对方欠我们的，应付＝我们欠对方的。当场结清（现结）的买卖也列出来：增加和收款/付款同时记，余额不变。',
        F_NOTE, align=AL, border=False)
    ws.row_dimensions[R_TO].height = 20
    ws.row_dimensions[R_CT].height = 18

    # ── 汇总 ──
    hfill = fill(C_WL)
    ws.merge_cells(f'A{R_SH}:E{R_SH}')
    put(ws, f'A{R_SH}', '汇　总', F_HDR, hfill, align=AC)
    for col in 'BCDE':
        ws[f'{col}{R_SH}'].border = BD
    for col, t in zip('FGHI', ('期初余额\n（起日以前）', '本期增加', '本期收款 / 付款', '期末余额')):
        put(ws, f'{col}{R_SH}', t, F_HDR, hfill, align=ACW)
    ws.merge_cells(f'J{R_SH}:{LAST}{R_SH}')
    put(ws, f'J{R_SH}', '说　明', F_HDR, hfill, align=AC)
    for col in 'KL':
        ws[f'{col}{R_SH}'].border = BD
    ws.row_dimensions[R_SH].height = 32
    g0 = lambda f: f'IF({U}="",0,ROUND({f},2))'
    for r, lbl, amt, chong, xjk, t in ((R_SAR, '应收（对方欠我们的）', '往_应收额', '收_冲应收', '现结收', '收款'),
                                         (R_SAP, '应付（我们欠对方的）', '往_应付额', '收_冲应付', '现结付', '付款')):
        ws.merge_cells(f'A{r}:E{r}')
        put(ws, f'A{r}', lbl, F_SUM_B, fill('FFFCE4D6'), align=AL)
        for col in 'BCDE':
            ws[f'{col}{r}'].border = BD
        wsum = lambda cond: f'SUMIFS({amt},往_往来单位,{EU},往_板块,{SG},{cond})'
        ssum = lambda cond: f'SUMIFS({chong},收_往来单位,{EU},收_板块,{SG},{cond})'
        vals = {
            'F': '=' + g0(f'{wsum(bef("往_日期"))}-{ssum(bef("收_日期"))}'),
            'G': '=' + g0(f'{wsum(per("往_日期"))}+{S(xjk)}'),
            'H': '=' + g0(f'{ssum(per("收_日期"))}+{S(xjk)}'),
            'I': f'=ROUND(F{r}+G{r}-H{r},2)',
        }
        for col, v in vals.items():
            put(ws, f'{col}{r}', v, F_SUM_V, None, MONEY, AR)
        ws.merge_cells(f'J{r}:{LAST}{r}')
        put(ws, f'J{r}', f'=IF({S(xjk)}=0,"","本期增加、{t}里都含现结 "&TEXT({S(xjk)},"#,##0.00")&" 元")', F_NOTE, None, align=AL)
        for col in 'KL':
            ws[f'{col}{r}'].border = BD
        ws.row_dimensions[r].height = 22
    r = R_SNET
    ws.merge_cells(f'A{r}:E{r}')
    put(ws, f'A{r}', '净额（应收－应付；对方欠我们为正）', F_SUM_B, fill('FFFCE4D6'), align=AL)
    for col in 'BCDE':
        ws[f'{col}{r}'].border = BD
    put(ws, f'F{r}', f'=ROUND(F{R_SAR}-F{R_SAP},2)', F_SUM_V, None, MONEY, AR)
    put(ws, f'G{r}', None, F_AUTO, FILL_AUTO)
    put(ws, f'H{r}', None, F_AUTO, FILL_AUTO)
    put(ws, f'I{r}', f'=ROUND(I{R_SAR}-I{R_SAP},2)', F_SUM_V, None, MONEY, AR)
    ws.merge_cells(f'J{r}:{LAST}{r}')
    put(ws, f'J{r}', (f'=IF({U}="","",IF(I{r}>0,"截至 "&TEXT({S1},"yyyy-mm-dd")&" 对方欠我们 "&TEXT(I{r},"#,##0.00")&" 元",'
                      f'IF(I{r}<0,"截至 "&TEXT({S1},"yyyy-mm-dd")&" 我们欠对方 "&TEXT(-I{r},"#,##0.00")&" 元","截至 "&TEXT({S1},"yyyy-mm-dd")&" 两清")))'),
        F_SUM_B, None, align=AL)
    for col in 'KL':
        ws[f'{col}{r}'].border = BD
    ws.row_dimensions[r].height = 22
    ws.row_dimensions[14].height = 10

    # ── 明细表头 ──
    heads = [('A', '日期'), ('B', '板块'), ('C', '业务内容（摘要）'), ('D', '数量\n（KG）'), ('E', '单价'), ('F', '应收增加'),
             ('G', '收款'), ('H', '应收余额'), ('I', '应付增加'), ('J', '付款'), ('K', '应付余额'), ('L', '备注')]
    header(ws, R_HDR, heads, C_WL, height=32)

    # ── 隐藏：排序键（_收 5000 条 ＋ _往 2000 条，键＝_收/_往 的排序键） ──
    for c, t in ((YN, '方块行号'), (YKS[0], '_收排序键'), (YKW[0], '_往排序键'), (HK, 'k'), (HC, '类型'), (HKEY, '键'), (HN, '第几条'), (HSRC, '哪张表'), (HUSE, '用途/类型')):
        ws[f'{c}{R0 - 1}'] = t
        ws[f'{c}{R0 - 1}'].font = F_HELP
    segok = lambda nm, n: f'OR({S("板块")}="",INDEX({nm},{n})={S("板块")})'
    for i in range(BLK):
        r = R0 + i
        ws[f'{YN}{r}'] = i + 1
        ws[f'{YN}{r}'].font = F_HELP
        for j, col in enumerate(YKS):
            n = f'(${YN}{r}+{j * BLK})'
            x = lambda k: f'INDEX(收_{k},{n})'
            ws[f'{col}{r}'] = (f'=IF({U}="","",IF(AND({x("有效")}=1,{x("往来单位")}={U},{x("内部转账")}=0,{x("日期")}>={S0},'
                               f'{x("日期")}<={S1},{segok("收_板块", n)}),{x("排序键")},""))')
        for j, col in enumerate(YKW):
            n = f'(${YN}{r}+{j * BLK})'
            x = lambda k: f'INDEX(往_{k},{n})'
            ws[f'{col}{r}'] = (f'=IF({U}="","",IF(AND({x("有效")}=1,{x("往来单位")}={U},{x("日期")}>={S0},'
                               f'{x("日期")}<={S1},{segok("往_板块", n)}),{x("排序键")},""))')
        for col in YKS + YKW:
            ws[f'{col}{r}'].font = F_HELP

    # ── 活动区：期初 → 明细 → 合计 → 共几笔 → 告知 → 签章 ──
    n = S('n')
    fm = {'A': DATE, 'D': '#,##0.##', 'E': '0.00##', 'F': MONEY, 'G': MONEY, 'H': MONEY, 'I': MONEY, 'J': MONEY, 'K': MONEY}
    al = {'A': AC, 'B': AC, 'C': AL, 'D': AR, 'E': AR, 'F': AR, 'G': AR, 'H': Alignment(horizontal='right', vertical='center', shrink_to_fit=True),
          'I': AR, 'J': AR, 'K': AR, 'L': AL}
    for kk in range(NREG):
        r = R0 + kk
        k = f'${HK}{r}'
        ws[f'{HK}{r}'] = kk + 1
        ws[f'{HC}{r}'] = (f'=IF({k}=1,"O",IF({k}<={n}+1,"D",IF({k}={n}+2,"T",IF({k}={n}+3,"C",IF({k}={n}+5,"F1",'
                          f'IF({k}={n}+6,"F2",IF({k}={n}+8,"G1",IF({k}={n}+11,"G2",IF({k}={n}+12,"G3","")))))))))')
        cd, key, ix_, src, use = (f'${c}{r}' for c in (HC, HKEY, HN, HSRC, HUSE))
        ws[f'{HKEY}{r}'] = f'=IF({cd}="D",SMALL({KEYS},{k}-1),0)'
        ws[f'{HN}{r}'] = f'=IF({key}=0,0,IF(MOD({key},10000)>5000,MOD({key},10000)-5000,MOD({key},10000)))'
        ws[f'{HSRC}{r}'] = f'=IF({key}=0,"",IF(MOD({key},10000)>5000,"W","S"))'
        ws[f'{HUSE}{r}'] = f'=IF({src}="W",INDEX(往_类型,{ix_}),IF({src}="S",INDEX(收_用途,{ix_}),""))'
        g = lambda f: f'IF({src}="W",INDEX(往_{f},{ix_}),INDEX(收_{f},{ix_}))'
        sv = lambda f: f'INDEX(收_{f},{ix_})'
        xjc = f'AND({use}<>"冲应收",{use}<>"冲应付")'
        nz = lambda x: f'IF({x}=0,"",{x})'
        prev = lambda c: f'{c}{r - 1}'
        above = lambda c: f'SUM({c}${R0}:{c}{r - 1})'      # 合计：只加它上面的行（不能把自己算进去）
        disp = {
            'A': {'D': g('日期')},
            'B': {'D': g('板块')},
            'C': {'O': f'IF({U}="","← 先在上面黄格选往来单位","期初余额（"&TEXT({S0},"yyyy-mm-dd")&" 以前）")',
                  'D': g('显示摘要'), 'T': '"本期合计"',
                  'C': f'IF({S("笔数")}>{NL},"⚠ 共 "&{S("笔数")}&" 笔，只列了前 {NL} 笔","共 "&{S("笔数")}&" 笔")',
                  'F1': '"以上往来如有不符，"', 'F2': '"请于收到之日起 7 日内告知。"',
                  'G1': '"本单位（盖章）："', 'G2': '"经办人："', 'G3': '"日期："'},
            'D': {'D': nz(g('数量'))},
            'E': {'D': nz(g('单价'))},
            'F': {'D': f'IF({src}="W",IF({use}="应收",INDEX(往_金额,{ix_}),""),IF({xjc},{nz(sv("收入"))},""))', 'T': above('F')},
            'G': {'D': f'IF({src}="W","",IF({use}="冲应收",{sv("冲应收")},IF({xjc},{nz(sv("收入"))},"")))', 'T': above('G')},
            'H': {'O': f'$F${R_SAR}', 'D': f'ROUND(N({prev("H")})+N(F{r})-N(G{r}),2)', 'T': prev('H'),
                  'G1': '"对方单位（盖章）："', 'G2': '"确认人："', 'G3': '"日期："'},
            'I': {'D': f'IF({src}="W",IF({use}="应付",INDEX(往_金额,{ix_}),""),IF({xjc},{nz(sv("支出"))},""))', 'T': above('I')},
            'J': {'D': f'IF({src}="W","",IF({use}="冲应付",{sv("冲应付")},IF({xjc},{nz(sv("支出"))},"")))', 'T': above('J')},
            'K': {'O': f'$F${R_SAP}', 'D': f'ROUND(N({prev("K")})+N(I{r})-N(J{r}),2)', 'T': prev('K')},
            'L': {'D': (f'IF({src}="W",IF(INDEX(往_金额,{ix_})<0,"折让","赊账"),IF({use}="冲应收","收回欠款",IF({use}="冲应付","支付欠款","现结")))'
                        f'&IF({g("经手人")}="","","·"&{g("经手人")})'),
                  'T': (f'IF({U}="","",IF({S("笔数")}>{NL},"⚠ 只列了前 {NL} 笔",IF({S("核对")}=1,"✓ 跟上面汇总一致","✗ 跟汇总对不上")))')},
        }
        if kk == 0:                       # 第一行只会是「期初」，不引用上一行
            for col in 'HK':
                disp[col].pop('D')
            for col in 'FGIJHK':
                disp[col].pop('T', None)
        for col, m in disp.items():
            cell = ws[f'{col}{r}']
            cell.value = _sw(cd, m)
            cell.font = F_TXT
            cell.alignment = al[col]
            if col in fm:
                cell.number_format = fm[col]
        for c in (HK, HC, HKEY, HN, HSRC, HUSE):
            ws[f'{c}{r}'].font = F_HELP
    reg = f'A{R0}:{LAST}{R1}'
    c0 = f'${HC}{R0}'
    rules = [
        (f'{c0}="O"', cf_fill('FFF2F2F2'), F_SUM_B, CF_BD),
        (f'{c0}="D"', None, None, CF_BD),
        (f'{c0}="T"', cf_fill('FFFCE4D6'), F_SUM_B, CF_BD),
        (f'{c0}="C"', None, F_GRAY, None),
        (f'OR({c0}="F1",{c0}="F2",{c0}="G1",{c0}="G2",{c0}="G3")', None, F_SIGN, None),
    ]
    for cond, fl, ft, bd in rules:
        kw = {}
        if fl is not None:
            kw['fill'] = fl
        if ft is not None:
            kw['font'] = ft
        if bd is not None:
            kw['border'] = bd
        ws.conditional_formatting.add(reg, FormulaRule(formula=[cond], stopIfTrue=True, **kw))
    ws.conditional_formatting.add(f'L{R0}:L{R1}', FormulaRule(formula=[f'AND({c0}="T",LEFT($L{R0},1)<>"✓")'], font=F_WARN))
    hide(ws, *[CL(i) for i in range(CI('N'), CI(YKW[-1]) + 1)])
    ws.freeze_panes = 'A5'
    print_setup(ws, f'{R_HDR}:{R_HDR}', landscape=True)
    ws.print_options.horizontalCentered = True
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q}!$A${R_HEAD}:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("最后一行")})')
