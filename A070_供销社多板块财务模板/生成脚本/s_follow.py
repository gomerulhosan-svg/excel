# -*- coding: utf-8 -*-
"""【应收应付跟进】（往来组）：截止日＝首页截止日 P_截止（只显示）。黄格：板块（空＝全部）、看哪类（全部/应收/应付）、只看没结清的（是/否）。
   ① 汇总：应收、应付各一行（累计发生、已收回/已付、余额、其中已逾期，按所选板块）＋「其中不在往来单位名单里的」差额行；
      右边账龄（全部板块）：未到期、逾期 1～30、31～90、90 天以上（按 往_未结、往_逾期天数）。
      ① 还有「7 天内到期」：约定日期在截止日～截止日＋7 天、还没结的（按所选板块）。
   ② 按往来单位：往来单位名单（第一次出现的）＋ 名单外的单位（应收应付登记里出现的、收支登记里收回/支付欠款出现的，最多 30 个），
      只列有往来的（再按「只看没结清」筛）；发生、收回、余额、已逾期按所选板块，约定日、逾期天数、最近日期按单位全部板块。合计行＝①。
      按紧急程度排（拿去催款）：已逾期（应收＋应付）从大到小 → 余额（应收＋应付）从大到小 → 名单顺序；
      排名＝1＋排在它前面的单位数（三个 COUNTIFS，精确比较，不用拼大数）。最后一列「跟进备注」＝【往来单位】的备注。
   ③ 逾期明细：往_未结>0 且逾期的每一笔（按约定日期从早到晚＝逾期最久的在前，最多 300 笔）。
   口径：设计.md 第 3 节（往_未结、往_逾期天数、往_最早未结、往_最新发生、收_最新收付 都按 P_截止 算）。
   ②③ 是一块「活动区」：隐藏列先算出每一行是什么（单位行 / 共几个 / ③标题 / ③表头 / ③明细 / 共几笔），③ 紧接在 ② 后面。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, Border, Side, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'V'
NL2, NL3, N_EXTRA = 300, 300, 30            # ② 最多列几个单位、③ 最多列几笔、名单外的单位最多收几个
NC = N_UNIT + N_EXTRA                       # 候选单位：名单 500 ＋ 名单外 30
R_SEC1, R_G1, R_H1, R_AR, R_ARX, R_AP, R_APX = 5, 6, 7, 8, 9, 10, 11
R_SEC2, R_G2, R_H2, R_TOT, R0 = 13, 14, 15, 16, 17
NREG = NL2 + NL3 + 5
R1 = R0 + NREG - 1

# ── 隐藏列 ──
HK, HC, HI = 'W', 'X', 'Y'                  # 活动区：第 k 行、这一行是什么、取第几个
HSR, HSP = 'Z', 'AA'                        # 活动区：状态的应收部分、应付部分（第 R0 行起）
SL, SV = 'Z', 'AA'                          # 标量：说明、值（第 3～15 行）
SC = ['板块', '板块条件', '看哪类', '看应收', '看应付', '只看没结清', '要列的单位数', 'n2', '逾期笔数', 'n3',
      '名单外单位数', '有往来的单位数', '最后一行']
SR = {k: 3 + i for i, k in enumerate(SC)}
CC = ['名称', '名单', '有效', 'i', '收发生', '收回', '收余额', '收逾期', '收约定', '收天数', '收最近',
      '付发生', '已付', '付余额', '付逾期', '付约定', '付天数', '付最近', '最近赊账', '往来',
      '序', '急', '余', '显示', '排名']                         # 序＝候选顺序（名单在前）；急＝已逾期合计；余＝余额合计
CCOL = {k: CL(CI('AC') + j) for j, k in enumerate(CC)}          # AC..BA
# 长的键列折成 500 行一列的方块（SMALL 可以对整块取第 k 小），隐藏列都在活动区的行数以内，打印、滚动都不会拖出几千行
BLK = 500
assert N_WL % BLK == 0 and N_CASH % BLK == 0 and BLK <= NREG
BN = CL(CI(CCOL['排名']) + 2)                                    # 方块的行号 1～500（第 n 条＝行号＋列偏移）（BC）
DKW = [CL(CI(BN) + 1 + j) for j in range(N_WL // BLK)]          # 名单外单位：_往 的键（BD..BG），键＝n
DKS = [CL(CI(DKW[-1]) + 1 + j) for j in range(N_CASH // BLK)]   # 名单外单位：_收 的键（BH..BQ），键＝10000＋n
OKS = [CL(CI(DKS[-1]) + 2 + j) for j in range(N_WL // BLK)]     # ③ 排序键（BS..BV）：约定日期×10000＋n
DK_RNG = f'${DKW[0]}${R0}:${DKS[-1]}${R0 + BLK - 1}'
OK_RNG = f'${OKS[0]}${R0}:${OKS[-1]}${R0 + BLK - 1}'

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_WHITE_B = Font(name=YH, sz=10, bold=True, color='FFFFFFFF')
F_ORANGE_B = Font(name=YH, sz=10, bold=True, color='FF833C0C')
F_GRAY = Font(name=YH, sz=9, color='FF808080')
F_OVER = Font(name=YH, sz=10, bold=True, color='FF9C0006')
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC00000')
C_GRP = 'FFED7D31'


def S(k):
    return f'${SV}${SR[k]}'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def _chain(pairs, default='""'):
    out = default
    for c, v in reversed(pairs):
        out = f'IF({c},{v},{out})'
    return out


def _sw(cd, m):
    return '=' + _chain([(f'{cd}="{k}"', v) for k, v in m.items()])


def _q(t):
    return '"' + t.replace('"', '""') + '"'


def crng(key):
    c = CCOL[key] if key in CCOL else key
    return f'${c}${R0}:${c}${R0 + NC - 1}'


TIP = ('💡 每个客户、供应商欠多少、欠了多久，全自动不用填。截止日期＝【首页】的截止日期（要换日期去首页改）。'
       '黄格：板块（空＝全部板块）、看哪类（全部 / 应收 / 应付）、只看没结清的（是＝余额不为 0 的才列；否＝有过往来的都列）。'
       '应收＝【应收应付登记】里赊出去的，减【收支登记】里「收回欠款」收回的；应付＝赊进来的，减「支付欠款」付掉的；先欠先还，'
       '过了约定日期还没结的就算逾期（逾期的行标红）。① 汇总和账龄；② 每个往来单位一行；③ 每一笔逾期没结的。'
       '②按紧急程度排：已逾期的多的在前，再按余额从大到小；最后一列「跟进备注」在【往来单位】表的备注里写（催过几次、答应哪天付）。'
       '当场结清（现结）的买卖不在这里，打印给对方核对用【对账单】。')


def build(wb, ctx=None):
    ws = wb[SH_FOLLOW]
    W = {'A': 5, 'B': 20, 'C': 12, 'D': 12.5, 'E': 16, 'F': 13.5, 'G': 13.5, 'H': 13.5, 'I': 13.5, 'J': 12.5, 'K': 8,
         'L': 12.5, 'M': 13.5, 'N': 13.5, 'O': 13.5, 'P': 13.5, 'Q': 12.5, 'R': 8, 'S': 12.5, 'T': 12.5, 'U': 32, 'V': 30}
    widths(ws, W)
    title(ws, '应收应付跟进', LAST, C_WL, TIP)
    wR, wP = S('看应收'), S('看应付')
    seg = S('板块条件')

    # ── 第 3、4 行：截止日（只显示）＋ 选择格 ──
    put(ws, 'A3', '截止日期', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.merge_cells('A3:B3')
    put(ws, 'C3', '=P_截止', F_AUTOB, FILL_AUTO, DATE, AC)
    c = put(ws, 'D3', '要换日期去【首页】改 →', Font(name=YH, sz=9, color='FF0563C1', underline='single'), align=AL, border=False)
    link(c, SH_HOME, HOME_END)
    ws.merge_cells('D3:E3')
    selector(ws, 'F3', '板块', 'G3', None, '=板块列表', prompt='空＝全部板块')
    selector(ws, 'H3', '看哪类', 'I3', '全部', '"全部,应收,应付"', prompt='全部 / 应收（别人欠我们的）/ 应付（我们欠别人的）')
    selector(ws, 'J3', '只看没结清的', 'L3', '是', '"是,否"', prompt='是＝余额不为 0 的才列；否＝有过往来的都列（包括已结清的）')
    ws.merge_cells('J3:K3')
    ws.merge_cells('M3:U3')
    put(ws, 'M3', f'=IF(AND({S("板块")}<>"",ISNA(MATCH({esc(S("板块"))},板块_名称,0))),"⚠ 板块「"&{S("板块")}&"」不在【基础资料】②里，下面都是 0",'
                  f'IF({S("名单外单位数")}>{N_EXTRA},"⚠ 不在往来单位名单里的单位超过 {N_EXTRA} 个，②只列了 {N_EXTRA} 个，请到【往来单位】补上",""))',
        F_WARN, align=AL, border=False)
    home_link(ws, 'V3')
    ws.merge_cells(f'A4:{LAST}4')
    put(ws, 'A4', (f'="现在看的：截止 "&TEXT(P_截止,"yyyy/mm/dd")&"；"&IF({S("板块")}="","全部板块",{S("板块")}&" 板块")&"；"'
                   f'&IF({S("看哪类")}="全部","应收和应付",{S("看哪类")})&"；"&IF({S("只看没结清")}=1,"只列没结清的往来单位","有过往来的单位都列")'
                   f'&"。（黄格空着＝全部板块；看哪类空着＝全部；只看没结清的空着＝是）"'), F_NOTE, align=AL, border=False)
    ws.row_dimensions[3].height = 22

    # ── 标量 ──
    sc = {
        '板块': '=TRIM(G3&"")',
        '板块条件': f'=IF({S("板块")}="","<>@@全部@@",{esc(S("板块"))})',
        '看哪类': '=IF(OR(TRIM(I3&"")="应收",TRIM(I3&"")="应付"),TRIM(I3&""),"全部")',
        '看应收': f'=IF({S("看哪类")}="应付",0,1)',
        '看应付': f'=IF({S("看哪类")}="应收",0,1)',
        '只看没结清': '=IF(TRIM(L3&"")="否",0,1)',
        '要列的单位数': f'=SUM({crng("显示")})',
        'n2': f'=MIN({S("要列的单位数")},{NL2})',
        '逾期笔数': f'=COUNT({OK_RNG})',
        'n3': f'=MIN({S("逾期笔数")},{NL3})',
        '名单外单位数': f'=COUNT({DK_RNG})',
        '有往来的单位数': f'=COUNTIFS({crng("往来")},">0")',
        '最后一行': f'={R0 - 1}+{S("n2")}+5+{S("n3")}',
    }
    for k in SC:
        ws[f'{SL}{SR[k]}'] = k
        ws[f'{SV}{SR[k]}'] = sc[k]
        ws[f'{SL}{SR[k]}'].font = ws[f'{SV}{SR[k]}'].font = F_HELP

    # ── ① 汇总 ＋ 账龄 ──
    section(ws, R_SEC1, 'A', LAST, f'="① 汇总（截至 "&TEXT(P_截止,"yyyy/mm/dd")&"）"', C_WL)
    lbl_fill = fill('FFD9E1F2')
    ws.merge_cells(f'A{R_G1}:E{R_H1}')
    put(ws, f'A{R_G1}', '', F_HDR, fill(C_WL), align=ACW)
    ws.merge_cells(f'F{R_G1}:J{R_G1}')
    put(ws, f'F{R_G1}', f'="按所选板块（"&IF({S("板块")}="","全部板块",{S("板块")})&"）"', F_HDR, fill(C_GRP), align=AC)
    ws.merge_cells(f'M{R_G1}:Q{R_G1}')
    put(ws, f'M{R_G1}', '账龄（全部板块；每笔没结的按逾期天数分）', F_HDR, fill(C_GRP), align=AC)
    ws.merge_cells(f'R{R_G1}:{LAST}{R_H1}')
    put(ws, f'R{R_G1}', '说明', F_HDR, fill(C_WL), align=AC)
    for col in 'GHIJNOPQSTUV':
        ws[f'{col}{R_G1}'].border = BD
    for col, t in zip('FGHIJ', ('累计发生', '已收回 / 已付', '余额', '其中已逾期', '7 天内到期\n（还没结的）')):
        put(ws, f'{col}{R_H1}', t, F_HDR, fill(C_WL), align=ACW)
    for col, t in zip('MNOPQ', ('未到期', '逾期 1～30 天', '逾期 31～90 天', '逾期 90 天以上', '合计\n（没结的）')):
        put(ws, f'{col}{R_H1}', t, F_HDR, fill(C_WL), align=ACW)
    ws.row_dimensions[R_H1].height = 32
    upto = '"<="&P_截止'
    lst_end = R0 + N_UNIT - 1
    blocks = [  # (行, 差额行, 类型, 往_额, 收_冲, 用途, 候选列前缀)
        (R_AR, R_ARX, '应收', '往_应收额', '收_冲应收', '冲应收', '收', '应收（别人欠我们的）', '多收了钱'),
        (R_AP, R_APX, '应付', '往_应付额', '收_冲应付', '冲应付', '付', '应付（我们欠别人的）', '多付了钱'),
    ]
    for r, rx, t, amt, chong, use, p, lbl, over in blocks:
        ws.merge_cells(f'A{r}:E{r}')
        put(ws, f'A{r}', lbl, F_KPI_L, lbl_fill, align=AL)
        ws.merge_cells(f'A{rx}:E{rx}')
        put(ws, f'A{rx}', '　其中：不在【往来单位】名单里的', F_TXT, None, align=AL)
        for col in 'BCDE':
            ws[f'{col}{r}'].border = ws[f'{col}{rx}'].border = BD
        f = {
            'F': f'=ROUND(SUMIFS({amt},往_日期,{upto},往_板块,{seg}),2)',
            'G': f'=ROUND(SUMIFS({chong},收_日期,{upto},收_板块,{seg}),2)',
            'H': f'=ROUND(F{r}-G{r},2)',
            'I': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_逾期天数,">0",往_板块,{seg}),2)',
            'J': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_约定日期,">="&P_截止,往_约定日期,"<="&(P_截止+7),往_板块,{seg}),2)',
            'M': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_逾期天数,0),2)',
            'N': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_逾期天数,">=1",往_逾期天数,"<=30"),2)',
            'O': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_逾期天数,">=31",往_逾期天数,"<=90"),2)',
            'P': f'=ROUND(SUMIFS(往_未结,往_类型,"{t}",往_逾期天数,">90"),2)',
            'Q': f'=ROUND(SUM(M{r}:P{r}),2)',
        }
        for col, v in f.items():
            put(ws, f'{col}{r}', v, F_AUTOB, None, MONEY, AR)
        keys = {'F': f'{p}发生', 'G': '收回' if p == '收' else '已付', 'H': f'{p}余额', 'I': f'{p}逾期'}
        for col, key in keys.items():
            c = CCOL[key]
            put(ws, f'{col}{rx}', f'=ROUND({col}{r}-SUM(${c}${R0}:${c}${lst_end}),2)', F_AUTO, None, MONEY, AR)
        for col in 'JMNOPQ':
            put(ws, f'{col}{rx}', None, F_AUTO, FILL_AUTO)
        bal = f'(SUMIFS({amt},往_日期,{upto})-SUMIFS({chong},收_日期,{upto}))'
        ws.merge_cells(f'R{r}:{LAST}{r}')
        put(ws, f'R{r}', (f'=IF(ROUND(Q{r}-{bal},2)=0,"✓ 没结的合计＝{t}余额（全部板块）",'
                          f'"没结的合计跟{t}余额（全部板块 "&TEXT({bal},"#,##0.00")&"）差 "&TEXT(Q{r}-{bal},"#,##0.00")'
                          f'&"：有单位{over}（看②的状态），或「{"收回欠款" if p == "收" else "支付欠款"}」没选往来单位")'), F_NOTE, None, align=ALW)
        ws.merge_cells(f'R{rx}:{LAST}{rx}')
        put(ws, f'R{rx}', (f'=IF(AND(F{rx}=0,G{rx}=0),"✓ 都在【往来单位】名单里",'
                           f'"⚠ 有不在名单里的单位（②里标「不在名单」），或「{"收回欠款" if p == "收" else "支付欠款"}」没选往来单位——到【数据校验】看")'),
            F_NOTE, None, align=ALW)
        for col in 'STUV':
            ws[f'{col}{r}'].border = ws[f'{col}{rx}'].border = BD
        ws.row_dimensions[r].height = 20
        ws.row_dimensions[rx].height = 20
    ws.row_dimensions[12].height = 8

    # ── ② 表头 ＋ 合计 ──
    section(ws, R_SEC2, 'A', LAST,
            (f'="② 按往来单位（截至 "&TEXT(P_截止,"yyyy/mm/dd")&"；"&IF({S("板块")}="","全部板块",{S("板块")}&" 板块")&"；"'
             f'&IF({S("看哪类")}="全部","应收和应付","只看"&{S("看哪类")})&"；"&IF({S("只看没结清")}=1,"只列没结清的","有过往来的都列")'
             f'&"）　逾期的行标红，下面接着是 ③ 逾期明细"'), C_WL)
    grp = [('A', 'E', '往来单位'), ('F', 'I', '应收（按所选板块）'), ('J', 'L', '应收（全部板块）'),
           ('M', 'P', '应付（按所选板块）'), ('Q', 'S', '应付（全部板块）'), ('T', 'T', '全部板块'), ('U', 'U', ''),
           ('V', 'V', '在【往来单位】表的备注里写')]
    for c1, c2, t in grp:
        if c1 != c2:
            ws.merge_cells(f'{c1}{R_G2}:{c2}{R_G2}')
        put(ws, f'{c1}{R_G2}', t, F_HDR, fill(C_GRP), align=AC)
        for i in range(CI(c1) + 1, CI(c2) + 1):
            put(ws, f'{CL(i)}{R_G2}', None, F_HDR, fill(C_GRP))
    ws.row_dimensions[R_G2].height = 18
    heads = [('A', '序号'), ('B', '往来单位'), ('C', '类型'), ('D', '联系人'), ('E', '电话'),
             ('F', '累计发生'), ('G', '已收回'), ('H', '余额'), ('I', '已逾期'), ('J', '最早没结的\n约定日'), ('K', '逾期\n天数'),
             ('L', '最近一次\n收款日'), ('M', '累计发生'), ('N', '已付'), ('O', '余额'), ('P', '已逾期'), ('Q', '最早没结的\n约定日'),
             ('R', '逾期\n天数'), ('S', '最近一次\n付款日'), ('T', '最近一笔\n赊账日'), ('U', '状态'), ('V', '跟进备注')]
    header(ws, R_H2, heads, C_WL, height=34)
    T = R_TOT
    put(ws, f'A{T}', None, F_AUTOB, FILL_TOT)
    put(ws, f'B{T}', '合计（＝①，含已结清的）', F_AUTOB, FILL_TOT, align=AL)
    ws.merge_cells(f'B{T}:E{T}')
    for col, src in zip('FGHI', 'FGHI'):
        put(ws, f'{col}{T}', f'=IF({wR}=1,{src}{R_AR},"")', F_AUTOB, FILL_TOT, MONEY, AR)
    for col, src in zip('MNOP', 'FGHI'):
        put(ws, f'{col}{T}', f'=IF({wP}=1,{src}{R_AP},"")', F_AUTOB, FILL_TOT, MONEY, AR)
    for col in 'CDEJKLQRSTV':
        put(ws, f'{col}{T}', None, F_AUTOB, FILL_TOT)
    put(ws, f'U{T}', f'="有往来的 "&{S("有往来的单位数")}&" 个，列出 "&{S("要列的单位数")}&" 个"', F_AUTOB, FILL_TOT, align=AL)
    ws.row_dimensions[T].height = 20

    # ── 隐藏：候选单位（名单 500 ＋ 名单外 30）各项指标 ──
    for k, c in CCOL.items():
        ws[f'{c}{R0 - 1}'] = k
        ws[f'{c}{R0 - 1}'].font = F_HELP
    for c, t in ((BN, '方块行号'), (DKW[0], '名单外:_往键'), (DKS[0], '名单外:_收键'), (OKS[0], '③:键'),
                 (HK, 'k'), (HC, '类型'), (HI, '第几个'), (HSR, '状态应收'), (HSP, '状态应付')):
        ws[f'{c}{R0 - 1}'] = t
        ws[f'{c}{R0 - 1}'].font = F_HELP
    for j in range(NC):
        r = R0 + j
        lst = j < N_UNIT
        g = lambda k: f'${CCOL[k]}{r}'
        nm, i_, v = g('名称'), g('i'), f'{g("有效")}=0'
        e = esc(nm)
        ws[g('名单').replace('$', '')] = 1 if lst else 0
        ws[g('i').replace('$', '')] = (j + 1) if lst else (j - N_UNIT + 1)
        ws[g('序').replace('$', '')] = j + 1
        f = {}
        if lst:
            f['名称'] = f'=INDEX(单位_名称,{i_})&""'
            f['有效'] = f'=IF({nm}="",0,IF(MATCH({e},单位_名称,0)={i_},1,0))'
        else:
            kd = f'SMALL({DK_RNG},{i_})'
            f['名称'] = (f'=IF({i_}>{S("名单外单位数")},"",IF({kd}<10000,INDEX(往_往来单位,{kd}),INDEX(收_往来单位,{kd}-10000))&"")')
            f['有效'] = f'=IF({nm}="",0,1)'
        for p, t, amt, chong, use in (('收', '应收', '往_应收额', '收_冲应收', '冲应收'), ('付', '应付', '往_应付额', '收_冲应付', '冲应付')):
            back = '收回' if p == '收' else '已付'
            f[f'{p}发生'] = f'=IF({v},0,SUMIFS({amt},往_往来单位,{e},往_日期,{upto},往_板块,{seg}))'
            f[back] = f'=IF({v},0,SUMIFS({chong},收_往来单位,{e},收_日期,{upto},收_板块,{seg}))'
            f[f'{p}余额'] = f'=ROUND({g(f"{p}发生")}-{g(back)},2)'
            f[f'{p}逾期'] = f'=IF({v},0,SUMIFS(往_未结,往_往来单位,{e},往_类型,"{t}",往_逾期天数,">0",往_板块,{seg}))'
            f[f'{p}约定'] = f'=IF({v},0,SUMIFS(往_约定日期,往_往来单位,{e},往_类型,"{t}",往_最早未结,1))'
            f[f'{p}天数'] = f'=IF({g(f"{p}约定")}=0,0,MAX(0,P_截止-{g(f"{p}约定")}))'
            f[f'{p}最近'] = f'=IF({v},0,SUMIFS(收_日期,收_往来单位,{e},收_用途,"{use}",收_最新收付,1))'
        f['最近赊账'] = (f'=IF({v},0,MAX({wR}*SUMIFS(往_日期,往_往来单位,{e},往_类型,"应收",往_最新发生,1),'
                       f'{wP}*SUMIFS(往_日期,往_往来单位,{e},往_类型,"应付",往_最新发生,1)))')
        cw = lambda t: f'COUNTIFS(往_往来单位,{e},往_类型,"{t}",往_有效,1,往_日期,{upto},往_板块,{seg})'
        cs = lambda u: f'COUNTIFS(收_往来单位,{e},收_用途,"{u}",收_有效,1,收_日期,{upto},收_板块,{seg})'
        f['往来'] = f'=IF({v},0,{wR}*({cw("应收")}+{cs("冲应收")})+{wP}*({cw("应付")}+{cs("冲应付")}))'
        # 要不要列、排第几：已逾期（应收＋应付）大的在前 → 余额（应收＋应付）大的在前 → 名单顺序
        f['急'] = f'=ROUND({wR}*{g("收逾期")}+{wP}*{g("付逾期")},2)'
        f['余'] = f'=ROUND({wR}*{g("收余额")}+{wP}*{g("付余额")},2)'
        f['显示'] = (f'=IF(AND({g("有效")}=1,{g("往来")}>0,OR({S("只看没结清")}=0,AND({wR}=1,{g("收余额")}<>0),'
                   f'AND({wP}=1,{g("付余额")}<>0))),1,0)')
        SH_, JI, YU, XU = crng('显示'), crng('急'), crng('余'), crng('序')
        f['排名'] = (f'=IF({g("显示")}=0,0,1+COUNTIFS({SH_},1,{JI},">"&{g("急")})'
                   f'+COUNTIFS({SH_},1,{JI},{g("急")},{YU},">"&{g("余")})'
                   f'+COUNTIFS({SH_},1,{JI},{g("急")},{YU},{g("余")},{XU},"<"&{g("序")}))')
        for k, val in f.items():
            ws[g(k).replace('$', '')] = val
        for k in CC:
            ws[g(k).replace('$', '')].font = F_HELP

    # ── 隐藏：名单外的单位（_往 的键＝n 在前、_收 的键＝10000＋n 在后；只取第一次出现的；_收 只看收回/支付欠款的行） ──
    #    ③ 逾期明细的排序键（约定日期×10000＋n）。都是 500 行一列的方块：第 n 条＝方块行号＋列偏移
    for i in range(BLK):
        r = R0 + i
        ws[f'{BN}{r}'] = i + 1
        ws[f'{BN}{r}'].font = F_HELP
        for j, col in enumerate(DKW):
            n = f'(${BN}{r}+{j * BLK})'
            u = f'INDEX(往_往来单位,{n})'
            ws[f'{col}{r}'] = (f'=IF({u}="","",IF(ISNA(MATCH({esc(u)},单位_名称,0)),'
                               f'IF(MATCH({esc(u)},往_往来单位,0)={n},{n},""),""))')
        for j, col in enumerate(DKS):
            n = f'(${BN}{r}+{j * BLK})'
            u = f'INDEX(收_往来单位,{n})'
            q = f'INDEX(收_用途,{n})'
            eu = esc(u)
            ws[f'{col}{r}'] = (f'=IF(OR({u}="",AND({q}<>"冲应收",{q}<>"冲应付")),"",'
                               f'IF(AND(ISNA(MATCH({eu},单位_名称,0)),ISNA(MATCH({eu},往_往来单位,0))),'
                               f'IF(COUNTIFS(收_往来单位,{eu},收_用途,"冲应收",收_n,"<"&{n})'
                               f'+COUNTIFS(收_往来单位,{eu},收_用途,"冲应付",收_n,"<"&{n})=0,10000+{n},""),""))')
        for j, col in enumerate(OKS):
            n = f'(${BN}{r}+{j * BLK})'
            x = lambda k: f'INDEX(往_{k},{n})'
            ws[f'{col}{r}'] = (f'=IF(AND({x("未结")}>0,{x("逾期天数")}>0,OR({S("看哪类")}="全部",{x("类型")}={S("看哪类")}),'
                               f'OR({S("板块")}="",{x("板块")}={S("板块")})),{x("约定日期")}*10000+{n},"")')
        for col in DKW + DKS + OKS:
            ws[f'{col}{r}'].font = F_HELP

    # ── 活动区：② 单位行 → 共几个 → 空行 → ③ 标题 → ③ 表头 → ③ 明细 → 共几笔 ──
    n2, n3 = S('n2'), S('n3')
    fm = {'A': '0', 'D': DATE, 'F': MONEY, 'G': MONEY, 'H': MONEY, 'I': MONEY, 'J': DATE, 'K': '0', 'L': DATE,
          'M': MONEY, 'N': MONEY, 'O': MONEY, 'P': MONEY, 'Q': DATE, 'R': '0', 'S': DATE, 'T': DATE}
    al = {c: AC for c in 'ACDJKLQRST'}
    al.update({c: AR for c in 'FGHIMNOP'})
    al.update({'B': AL, 'E': AL, 'U': AL, 'V': AL})
    for kk in range(NREG):
        r = R0 + kk
        k = f'${HK}{r}'
        ws[f'{HK}{r}'] = kk + 1
        ws[f'{HC}{r}'] = (f'=IF({k}<={n2},"U",IF({k}={n2}+1,"N2",IF({k}={n2}+3,"S3",IF({k}={n2}+4,"H3",'
                          f'IF(AND({k}>{n2}+4,{k}<={n2}+4+{n3}),"D3",IF({k}={n2}+5+{n3},"N3",""))))))')
        cd = f'${HC}{r}'
        ix = f'${HI}{r}'
        ws[f'{HI}{r}'] = (f'=IF({cd}="U",IFERROR(MATCH({k},{crng("排名")},0),0),'
                          f'IF({cd}="D3",IF({k}-{n2}-4>{S("逾期笔数")},0,MOD(SMALL({OK_RNG},{k}-{n2}-4),10000)),0))')
        cv = lambda key: f'INDEX({crng(key)},{ix})'
        w = lambda key: f'INDEX(往_{key},{ix})'
        inl = f'{cv("名单")}=1'
        ui = cv('i')
        date_or_blank = lambda key, want: f'IF(AND({want}=1,{cv(key)}>0),{cv(key)},"")'
        # 状态（应收部分、应付部分）
        for col, want, b, o, d, more, t in ((HSR, wR, 'H', 'I', 'K', '多收了 ', '应收'), (HSP, wP, 'O', 'P', 'R', '多付了 ', '应付')):
            ws[f'{col}{r}'] = (f'=IF(OR({cd}<>"U",{want}=0),"",IF(N(${b}{r})<0,"{more}"&TEXT(-N(${b}{r}),"#,##0.00"),'
                               f'IF(N(${b}{r})=0,"",IF(N(${o}{r})>0,"{t}逾期"&IF(N(${d}{r})>0," "&N(${d}{r})&" 天",""),"{t}未到期"))))')
        sr, sp = f'${HSR}{r}', f'${HSP}{r}'
        disp = {
            'A': {'U': k, 'H3': '"序号"', 'D3': f'{k}-{n2}-4'},
            'B': {'U': cv('名称'), 'N2': f'IF({n2}=0,"没有要列的往来单位","共 "&{S("要列的单位数")}&" 个往来单位")',
                  'S3': '"③ 逾期没结的明细"', 'H3': '"往来单位"', 'D3': w('往来单位'),
                  'N3': f'IF({S("逾期笔数")}=0,"没有逾期没结的 ✓","共 "&{S("逾期笔数")}&" 笔逾期没结的")'},
            'C': {'U': f'IF({inl},IF(INDEX(单位_类型,{ui})="","（没填类型）",INDEX(单位_类型,{ui})),"⚠ 不在名单")',
                  'H3': '"应收/应付"', 'D3': w('类型'), 'S3': f'"共 "&{S("逾期笔数")}&" 笔"'},
            'D': {'U': f'IF({inl},INDEX(单位_联系人,{ui}),"")', 'H3': '"日期"', 'D3': w('日期')},
            'E': {'U': f'IF({inl},INDEX(单位_电话,{ui}),"")', 'H3': '"摘要"', 'D3': w('显示摘要'),
                  'S3': '"逾期最久的在前"', 'N3': f'IF({S("逾期笔数")}>{NL3},"⚠ 只列了前 {NL3} 笔","")'},
            'F': {'U': f'IF({wR}=1,{cv("收发生")},"")', 'H3': '"金额"', 'D3': w('金额')},
            'G': {'U': f'IF({wR}=1,{cv("收回")},"")', 'H3': '"还没结的"', 'D3': w('未结')},
            'H': {'U': f'IF({wR}=1,{cv("收余额")},"")', 'H3': '"约定日期"', 'D3': f'TEXT({w("约定日期")},"yyyy/mm/dd")'},
            'I': {'U': f'IF({wR}=1,{cv("收逾期")},"")', 'H3': '"逾期天数"', 'D3': f'{w("逾期天数")}&" 天"'},
            'J': {'U': date_or_blank('收约定', wR), 'H3': '"板块"', 'D3': w('板块')},
            'K': {'U': f'IF(AND({wR}=1,{cv("收天数")}>0),{cv("收天数")},"")', 'H3': '"经手人"', 'D3': w('经手人')},
            'L': {'U': date_or_blank('收最近', wR), 'H3': '"登记行号"', 'D3': f'"第 "&{w("录入行")}&" 行"'},
            'M': {'U': f'IF({wP}=1,{cv("付发生")},"")'},
            'N': {'U': f'IF({wP}=1,{cv("已付")},"")'},
            'O': {'U': f'IF({wP}=1,{cv("付余额")},"")'},
            'P': {'U': f'IF({wP}=1,{cv("付逾期")},"")'},
            'Q': {'U': date_or_blank('付约定', wP)},
            'R': {'U': f'IF(AND({wP}=1,{cv("付天数")}>0),{cv("付天数")},"")'},
            'S': {'U': date_or_blank('付最近', wP)},
            'T': {'U': f'IF({cv("最近赊账")}>0,{cv("最近赊账")},"")'},
            'U': {'U': f'IF({sr}&{sp}="","已结清",{sr}&IF(AND({sr}<>"",{sp}<>""),"；","")&{sp})',
                  'N2': (f'IF({S("要列的单位数")}>{NL2},"⚠ 超过 {NL2} 个，只列了前 {NL2} 个",'
                         f'IF(AND({S("只看没结清")}=1,{S("有往来的单位数")}>{S("要列的单位数")}),'
                         f'"另有 "&({S("有往来的单位数")}-{S("要列的单位数")})&" 个已结清的没列出",""))')},
            'V': {'U': f'IF({inl},INDEX(单位_备注,{ui}),"")'},
        }
        for col, m in disp.items():
            cell = ws[f'{col}{r}']
            cell.value = _sw(cd, m)
            cell.font = F_TXT
            cell.alignment = al.get(col, AC)
            if col in fm:
                cell.number_format = fm[col]
        for col in (HK, HC, HI, HSR, HSP):
            ws[f'{col}{r}'].font = F_HELP
    reg = f'A{R0}:{LAST}{R1}'
    c0 = f'${HC}{R0}'
    ws.conditional_formatting.add(f'I{R0}:I{R1}', FormulaRule(formula=[f'{c0}="D3"'], font=F_WARN, border=CF_BD, stopIfTrue=True))
    reg3 = f'A{R0}:L{R1}'                  # ③ 只用 A～L
    rules = [
        (reg, f'AND({c0}="U",OR(N($I{R0})>0,N($P{R0})>0))', cf_fill('FFFFC7CE'), F_OVER, CF_BD),
        (reg, f'{c0}="U"', None, None, CF_BD),
        (reg3, f'{c0}="H3"', cf_fill(C_WL), F_WHITE_B, CF_BD),
        (reg3, f'{c0}="S3"', cf_fill('FFFCE4D6'), F_ORANGE_B, None),
        (reg3, f'{c0}="D3"', None, None, CF_BD),
        (reg, f'OR({c0}="N2",{c0}="N3")', None, F_GRAY, None),
    ]
    for rg, cond, fl, ft, bd in rules:
        kw = {}
        if fl is not None:
            kw['fill'] = fl
        if ft is not None:
            kw['font'] = ft
        if bd is not None:
            kw['border'] = bd
        ws.conditional_formatting.add(rg, FormulaRule(formula=[cond], stopIfTrue=True, **kw))

    hide(ws, *[CL(i) for i in range(CI(HK), CI(OKS[-1]) + 1)])
    ws.freeze_panes = 'C5'
    print_setup(ws, f'{R_G2}:{R_H2}', landscape=True)          # 每页重复 ② 的表头
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${LAST}$1:${LAST}${R1},{q}!{S("最后一行")})')
