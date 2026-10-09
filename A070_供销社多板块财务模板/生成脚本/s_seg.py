# -*- coding: utf-8 -*-
"""板块归档表（绿）：培训学校、为农服务中心、粮食烘干、粮食购销、供销社（layout.SEGMENTS，表名固定）。
   每张表＝从【收支登记】【应收应付登记】自动拆出来的这个板块的台账，保持用户原表的列（左边、顺序不变），右边加补充列：
     通用（培训学校、为农服务中心、供销社）：日期｜摘要｜收入｜支出｜结余｜经手人
     粮食烘干：日期｜摘要｜客户名称｜烘干重量KG｜烘干单价｜收入｜支出｜结余｜经手人
     粮食购销：日期｜摘要｜客户名称｜购进数量KG｜购进单价｜销售数量｜销售单价｜粮食存量｜收入｜支出｜结余｜经手人
     右边：收支项目｜往来单位（只通用表）｜应收（赊）｜应付（赊）｜账户｜备注｜来源
   黄格：板块（默认＝表名）、起（空＝建账日）、止（空＝首页截止日）、经手人（空＝全部）。
   口径见 设计.md §3：结余用 收_板块收入/收_板块支出（不含内部转账），存量两张表都算；经手人只筛「本期」的数和清单，余额不按经手人分。

   清单的排法（省公式）：5 张表共用两列「分组排序键」（隐藏表 _序 的 A、B 列，第 2 行起：_收 5000 条＋_往 2000 条；C 列是第几条）：
     A ＝ 板块序号×1E9 ＋ 排序键（排序键＝日期×10000＋n，_往 再加 5000，同一天收支在前；内部转账、没板块的为空）
     B ＝ (板块序号×(N_PER+1)＋经手人序号)×1E9 ＋ 排序键
   这样「某板块、某段日期（、某经手人）」在这一列里正好是连续的一段：起点＝COUNTIF(列,"<"&下限)，条数＝COUNTIFS(列,">="&下限,列,"<="&上限)，
   第 k 条＝MOD(SMALL(列,起点+k),10000)（≤5000 是 _收 第几条，>5000 是 _往 第 几−5000 条）。每张表只要 1000 行显示公式，不用各自 7000 行排序键。
"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font
from layout import *
from common import *

# ───────────── 外观 ─────────────
GREEN_H = 'FF70AD47'                                  # 补充列表头（比用户原表的列浅一点）
F_SUMV = Font(name=YH, sz=11, bold=True, color='FFC00000')
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'                # 零不显示（清单用）
QTY = '#,##0.##'
QTY_B = '#,##0.##;-#,##0.##;""'
PRICE_B = '#,##0.00##;-#,##0.00##;""'
FILL_SHE = 'FFFDE9D9'                                 # 赊账（应收应付登记）的行：淡橙

# ───────────── 行（5 张表一样） ─────────────
SEL = dict(板块='B3', 起='E3', 止='G3', 经手人='I3')   # 黄格
R_SUM_L, R_SUM_V, R_NOTE, HDR, R_TOT, R_OPEN, R0 = 5, 6, 7, 8, 9, 10, 11
SHOW = 1000                                           # 清单最多显示几条
R_END = R0 + SHOW - 1

# ───────────── 列（用户原表的列在前、顺序不变；右边补充列） ─────────────
_USER = {
    '通用': ['日期', '摘要', '收入', '支出', '结余', '经手人'],
    '烘干': ['日期', '摘要', '往来单位', '烘干重量', '烘干单价', '收入', '支出', '结余', '经手人'],
    '购销': ['日期', '摘要', '往来单位', '购进数量', '购进单价', '销售数量', '销售单价', '存量', '收入', '支出', '结余', '经手人'],
}
_EXTRA = {
    '通用': ['收支项目', '往来单位', '应收', '应付', '账户', '备注', '来源'],
    '烘干': ['收支项目', '应收', '应付', '账户', '备注', '来源'],
    '购销': ['收支项目', '应收', '应付', '账户', '备注', '来源'],
}
COLS = {k: {f: CL(i + 1) for i, f in enumerate(_USER[k] + _EXTRA[k])} for k in _USER}   # COLS[kind]['结余'] → 列字母
LAST = {k: CL(len(_USER[k]) + len(_EXTRA[k])) for k in _USER}
HEAD = {'日期': '日期', '摘要': '摘要', '收入': '收入', '支出': '支出', '结余': '结余', '经手人': '经手人',
        '烘干重量': '烘干重量KG', '烘干单价': '烘干单价', '购进数量': '购进数量KG', '购进单价': '购进单价',
        '销售数量': '销售数量', '销售单价': '销售单价', '存量': '粮食存量', '收支项目': '收支项目',
        '应收': '应收（赊）\n收回为负', '应付': '应付（赊）\n付出为负', '账户': '账户', '备注': '备注',
        '来源': '来源\n（去这一行改）'}
WIDTH = {'日期': 11, '摘要': 28, '往来单位': 13, '收入': 13, '支出': 13, '结余': 14, '经手人': 10, '烘干重量': 12,
         '烘干单价': 11, '购进数量': 12, '购进单价': 11, '销售数量': 11, '销售单价': 11, '存量': 12, '收支项目': 12,
         '应收': 13, '应付': 13, '账户': 10, '备注': 24, '来源': 22}

# ───────────── 隐藏列 ─────────────
KC, ZC, SC_COL, SC_LBL = 'Y', 'Z', 'AA', 'AB'          # 第几条(常数)、对应 _收/_往 第几条、标量、标量说明
HK_SHEET = H_SORT                                      # 共用的分组排序键放在隐藏表 _序（删哪张板块表都不影响别的）
HK1C, HK2C, HN, HK_R0 = 'A', 'B', 'C', 2               # 键 1（板块）、键 2（板块＋经手人）、n（常数，辅助）
N_HK = N_CASH + N_WL
HK1 = f'{HK_SHEET}!${HK1C}${HK_R0}:${HK1C}${HK_R0 + N_HK - 1}'
HK2 = f'{HK_SHEET}!${HK2C}${HK_R0}:${HK2C}${HK_R0 + N_HK - 1}'
PER_W = N_PER + 1

# 标量（隐藏 AA 列第几行）——首页、查询要链过来可以用 key_cell(名)
SC = dict(板块=1, 板块条件=2, 板块号=3, 起=4, 止=5, 经手人=6, 经手人条件=7, 经手人号=8, 下限=9, 上限=10, 起点=11, 条数=12,
          期初结余=13, 本期收入=14, 本期支出=15, 期末结余=16, 应收余额=17, 应付余额=18, 本期应收增减=19, 本期应付增减=20,
          期初存量=21, 本期购进=22, 本期销售=23, 期末存量=24, 烘干重量=25, 烘干收入=26)


def key_cell(name):
    return f'${SC_COL}${SC[name]}'


S = key_cell

# 顶部汇总（第 5 行标签、第 6 行数）：(标量名, 标签, 格式)
_SUM_BASE = [('期初结余', '期初结余\n（起日前一天）', MONEY), ('本期收入', '本期收入', MONEY), ('本期支出', '本期支出', MONEY),
             ('期末结余', '期末结余\n（到止）', MONEY), ('应收余额', '应收余额（到止）\n人家欠的', MONEY),
             ('应付余额', '应付余额（到止）\n欠人家的', MONEY)]
SUM_ITEMS = {
    '通用': _SUM_BASE,
    '烘干': _SUM_BASE + [('烘干重量', '本期烘干重量KG\n（收入行的数量）', QTY), ('烘干收入', '烘干收入\n（含赊账）', MONEY)],
    '购销': _SUM_BASE + [('期初存量', '期初存量KG\n（起日前一天）', QTY), ('本期购进', '本期购进KG\n（含赊购）', QTY),
                       ('本期销售', '本期销售KG\n（含赊销）', QTY), ('期末存量', '期末存量KG\n（到止）', QTY)],
}
SUMMARY = {k: {nm: f'{CL(2 + i)}{R_SUM_V}' for i, (nm, _l, _f) in enumerate(v)} for k, v in SUM_ITEMS.items()}


def g(f, z, wf=None):
    """第 z 条：z≤N_CASH 取 _收 第 z 条，否则取 _往 第 z−N_CASH 条"""
    return f'IF({z}<={N_CASH},INDEX(收_{f},{z}),INDEX(往_{wf or f},{z}-{N_CASH}))'


def per(t, a=None, b=None):
    a, b = a or S('起'), b or S('止')
    return dr(f'{t}_日期', a, b)


# ═════════════════════════════ 共用的分组排序键（隐藏表 _序） ═════════════════════════════
def build_keys(ws):
    ws[f'{HK1C}1'] = '板块序号×1E9＋排序键（板块表清单用）'
    ws[f'{HK2C}1'] = f'(板块序号×{PER_W}＋经手人序号)×1E9＋排序键'
    ws[f'{HN}1'] = f'第几条（_收 1～{N_CASH}，再 _往 1～{N_WL}）'
    for c in (HN, HK1C, HK2C):
        ws[f'{c}1'].font = F_HELP
    for i in range(N_HK):
        r = HK_R0 + i
        t = '收' if i < N_CASH else '往'
        ws[f'{HN}{r}'] = i + 1 if i < N_CASH else i - N_CASH + 1
        n = f'${HN}{r}'
        skip = f'INDEX(收_内部转账,{n})=1,' if t == '收' else ''
        ws[f'{HK1C}{r}'] = (f'=IF(INDEX({t}_排序键,{n})="","",IF(OR({skip}INDEX({t}_板块,{n})=""),"",'
                            f'IFERROR(MATCH(INDEX({t}_板块,{n}),板块_名称,0)*1E9+INDEX({t}_排序键,{n}),"")))')
        ws[f'{HK2C}{r}'] = (f'=IF(${HK1C}{r}="","",(INT(${HK1C}{r}/1E9)*{PER_W}+IF(INDEX({t}_经手人,{n})="",0,'
                            f'IFERROR(MATCH(INDEX({t}_经手人,{n}),经手人_姓名,0),0)))*1E9+MOD(${HK1C}{r},1E9))')
        for c in (HN, HK1C, HK2C):
            ws[f'{c}{r}'].font = F_HELP


# ═════════════════════════════ 一张板块表 ═════════════════════════════
def _scalars(ws, seg):
    a, b, pe, ec = S('起'), S('止'), S('板块条件'), S('经手人条件')
    lt, le = f'"<"&{a}', f'"<="&{b}'
    code = f'IF({S("经手人")}="",{S("板块号")},{S("板块号")}*{PER_W}+{S("经手人号")})'

    def cnt_hk(crit):
        return f'IF({S("经手人")}="",{crit.format(HK=HK1)},{crit.format(HK=HK2)})'

    def stock(op, d):
        q = f'"{op}"&{d}'
        return (f'SUMIFS(收_入库量,收_板块,{pe},收_日期,{q})+SUMIFS(往_入库量,往_板块,{pe},往_日期,{q})'
                f'-SUMIFS(收_出库量,收_板块,{pe},收_日期,{q})-SUMIFS(往_出库量,往_板块,{pe},往_日期,{q})')

    def bal(op, d):
        q = f'"{op}"&{d}'
        return (f'SUMIFS(板块_期初结余,板块_名称,{pe})+SUMIFS(收_板块收入,收_板块,{pe},收_日期,{q})'
                f'-SUMIFS(收_板块支出,收_板块,{pe},收_日期,{q})')

    def rcv(amt, chong, op, d):
        q = f'"{op}"&{d}'
        return f'SUMIFS(往_{amt},往_板块,{pe},往_日期,{q})-SUMIFS(收_{chong},收_板块,{pe},收_日期,{q})'

    pp = lambda t: f'{t}_板块,{pe},{per(t)},{t}_经手人,{ec}'          # 本期、按经手人
    inc_s = '收_方向,"收入",收_用途,"<>冲应收",收_用途,"<>内部转账"'     # 收入方向的收支行（烘干重量、烘干收入用；不认项目名）
    f = {
        '板块': f'=IF(TRIM({SEL["板块"]}&"")="","{seg}",TRIM({SEL["板块"]}&""))',
        '板块条件': f'={esc(S("板块"))}',
        '板块号': f'=IFERROR(MATCH({S("板块")},板块_名称,0),0)',
        '起': f'=IF(ISNUMBER({SEL["起"]}),INT({SEL["起"]}),P_建账日)',
        '止': f'=IF(ISNUMBER({SEL["止"]}),INT({SEL["止"]}),P_截止)',
        '经手人': f'=TRIM({SEL["经手人"]}&"")',
        '经手人条件': f'=IF({S("经手人")}="","<>@@全部@@",{esc(S("经手人"))})',
        '经手人号': f'=IF({S("经手人")}="",0,IFERROR(MATCH({S("经手人")},经手人_姓名,0),-1))',
        '下限': f'={code}*1E9+{a}*10000',
        '上限': f'={code}*1E9+{b}*10000+9999',
        '起点': '=' + cnt_hk(f'COUNTIF({{HK}},"<"&TEXT({S("下限")},"0"))'),
        '条数': (f'=IF(OR({S("板块号")}=0,{S("经手人号")}<0,{a}>{b}),0,'
               + cnt_hk(f'COUNTIFS({{HK}},">="&TEXT({S("下限")},"0"),{{HK}},"<="&TEXT({S("上限")},"0"))') + ')'),
        '期初结余': f'=ROUND({bal("<", a)},2)',
        '本期收入': f'=ROUND(SUMIFS(收_板块收入,{pp("收")}),2)',
        '本期支出': f'=ROUND(SUMIFS(收_板块支出,{pp("收")}),2)',
        '期末结余': f'=ROUND({bal("<=", b)},2)',
        '应收余额': f'=ROUND({rcv("应收额", "冲应收", "<=", b)},2)',
        '应付余额': f'=ROUND({rcv("应付额", "冲应付", "<=", b)},2)',
        '本期应收增减': f'=ROUND(SUMIFS(往_应收额,{pp("往")})-SUMIFS(收_冲应收,{pp("收")}),2)',
        '本期应付增减': f'=ROUND(SUMIFS(往_应付额,{pp("往")})-SUMIFS(收_冲应付,{pp("收")}),2)',
        '期初存量': f'=ROUND(SUMIFS(板块_期初存量,板块_名称,{pe})+{stock("<", a)},3)',
        '本期购进': f'=ROUND(SUMIFS(收_入库量,{pp("收")})+SUMIFS(往_入库量,{pp("往")}),3)',
        '本期销售': f'=ROUND(SUMIFS(收_出库量,{pp("收")})+SUMIFS(往_出库量,{pp("往")}),3)',
        '期末存量': f'=ROUND(SUMIFS(板块_期初存量,板块_名称,{pe})+{stock("<=", b)},3)',
        '烘干重量': (f'=ROUND(SUMIFS(收_数量,{inc_s},收_有效,1,{pp("收")})'
                 f'+SUMIFS(往_数量,往_类型,"应收",往_有效,1,{pp("往")}),3)'),
        '烘干收入': (f'=ROUND(SUMIFS(收_板块收入,{inc_s},{pp("收")})'
                 f'+SUMIFS(往_应收额,往_类型,"应收",{pp("往")}),2)'),
    }
    for k, r in SC.items():
        ws[f'{SC_COL}{r}'] = f[k]
        ws[f'{SC_LBL}{r}'] = k
        ws[f'{SC_COL}{r}'].font = ws[f'{SC_LBL}{r}'].font = F_HELP


def _tip(kind):
    t = ('💡 这张表不用填：【收支登记】【应收应付登记】里业务板块选了这个板块的每一笔，自动按日期列在这里（同一天先列收支、再列赊账）。'
         '黄格可以改：板块（默认就是本表；基础资料里改了板块名，在这里重选）、起止日期（空＝建账日～首页截止日）、经手人（空＝全部）。'
         '结余只按收支登记的钱滚动（内部转账不算）；淡橙色的行是赊账（应收应付登记），只记在「应收（赊）」「应付（赊）」，不动结余；'
         '收回欠款、支付欠款在这两列显示负数。选了经手人：只列他经手的，本期数只算他的，结余、余额还是整个板块的。'
         '要改哪一笔，按最后一列「来源」去登记表改。')
    if kind == '购销':
        t += '粮食存量＝期初存量＋购进−销售（赊购、赊销的也算）；不是购进、销售的行有数量的，写在摘要后面（数量×单价）。'
    elif kind == '烘干':
        t += ('烘干重量、烘干单价就是登记表里的数量、单价；只填了数量、单价的，收入按数量×单价算。'
              '上面的「烘干重量」「烘干收入」＝收入类的行（不含收回欠款）＋赊出去的应收，不管收支项目叫什么。')
    else:
        t += '登记表里填了数量的，写在摘要后面（数量×单价）。'
    return t


def build_seg(ws, seg, kind):
    C = COLS[kind]
    last = LAST[kind]
    widths(ws, {c: WIDTH.get(f, 12) for f, c in C.items()})
    title(ws, f'={S("板块")}&" · 收支台账（从两张登记表自动拆分）"', last, C_SEG, _tip(kind))

    # ── 选择格 ──
    selector(ws, 'A3', '板块', SEL['板块'], seg, '=板块列表',
             prompt='默认＝本表；从下拉选别的板块也行（基础资料里改了板块名，就在这里重选）')
    selector(ws, 'D3', '起', SEL['起'], None, fmt=DATE, prompt='空＝建账日')
    selector(ws, 'F3', '止', SEL['止'], None, fmt=DATE, prompt='空＝首页截止日')
    selector(ws, 'H3', '经手人', SEL['经手人'], None, '=经手人列表', prompt='空＝全部经手人')
    dv_date(ws, SEL['起'])
    dv_date(ws, SEL['止'])
    home_link(ws, f'{last}3')
    put(ws, 'A4', '实际用的', F_NOTE, align=AC)
    put(ws, 'B4', f'={S("板块")}&IF({S("板块号")}=0,"（⚠ 不在【基础资料】②里）","")', F_AUTOB, FILL_AUTO, align=AC)
    put(ws, 'D4', '空＝建账日→', F_NOTE, align=AR)
    put(ws, 'E4', f'={S("起")}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'F4', '空＝截止日→', F_NOTE, align=AR)
    put(ws, 'G4', f'={S("止")}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'H4', '空＝全部→', F_NOTE, align=AR)
    put(ws, 'I4', f'=IF({S("经手人")}="","全部",{S("经手人")})', F_AUTOB, FILL_AUTO, align=AC)
    ws.row_dimensions[3].height = 22

    _scalars(ws, seg)

    # ── 顶部汇总 ──
    ws.merge_cells(f'A{R_SUM_L}:A{R_SUM_V}')
    put(ws, f'A{R_SUM_L}', '本期汇总', F_KPI_L, fill('FFD9E1F2'), align=ACW)
    for i, (nm, lbl, fmt) in enumerate(SUM_ITEMS[kind]):
        col = CL(2 + i)
        put(ws, f'{col}{R_SUM_L}', lbl, F_KPI_L, fill('FFD9E1F2'), align=ACW)
        put(ws, f'{col}{R_SUM_V}', f'={S(nm)}', F_SUMV, FILL_AUTO, fmt, AR)
    ws.row_dimensions[R_SUM_L].height = 32
    ws.row_dimensions[R_SUM_V].height = 22

    # ── 清单 ──
    n = S('条数')
    z_rng = lambda col: f'${col}${R0}:${col}${R_END}'
    has_stock = '存量' in C
    lastbal = f'INDEX({z_rng(C["结余"])},{n})'
    ok_tail = (f'IF(ROUND({lastbal}-{S("期末结余")},2)=0,"；最后一行结余＝期末结余 ✓","；⚠ 最后一行结余跟期末结余对不上")')
    if has_stock:
        ok_tail += (f'&IF(ROUND(INDEX({z_rng(C["存量"])},{n})-{S("期末存量")},3)=0,"，存量＝期末存量 ✓",'
                    f'"；⚠ 最后一行存量跟期末存量对不上")')
    ws.merge_cells(f'A{R_NOTE}:{last}{R_NOTE}')
    put(ws, f'A{R_NOTE}', (
        f'=IF({S("起")}>{S("止")},"⚠ 起晚于止，请改黄格里的日期",'
        f'IF({S("板块号")}=0,"⚠ 板块「"&{S("板块")}&"」不在【基础资料】②业务板块里，请在黄格重选",'
        f'IF({S("经手人号")}<0,"⚠ 经手人「"&{S("经手人")}&"」不在【基础资料】⑤里，清单列不出来（上面的本期数照算）",'
        f'IF({n}=0,"这段时间没有流水",'
        f'IF({n}>{SHOW},"⚠ 共 "&{n}&" 条，只显示前 {SHOW} 条，请把「起」改晚一点（上面的汇总是全部的）",'
        f'"共 "&{n}&" 条，按日期排（同一天先列收支、再列赊账）"&IF({S("经手人")}="",{ok_tail},'
        f'"；只列「"&{S("经手人")}&"」经手的，结余、存量是整个板块的"))))))'), F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'A{R_NOTE}', FormulaRule(formula=[f'LEFT($A${R_NOTE},1)="⚠"'], font=F_RED))
    ws.row_dimensions[R_NOTE].height = 20

    header(ws, HDR, [(c, HEAD[f] if f != '往来单位' else ('往来单位' if kind == '通用' else '客户名称'))
                     for f, c in C.items() if f in _USER[kind]], C_SEG)
    header(ws, HDR, [(c, HEAD.get(f, f)) for f, c in C.items() if f not in _USER[kind]], GREEN_H)
    ws.row_dimensions[HDR].height = 34

    fmts = {'日期': DATE, '收入': MONEY_B, '支出': MONEY_B, '结余': MONEY, '应收': MONEY_B, '应付': MONEY_B,
            '烘干重量': QTY_B, '烘干单价': PRICE_B, '购进数量': QTY_B, '购进单价': PRICE_B, '销售数量': QTY_B,
            '销售单价': PRICE_B, '存量': QTY}
    aligns = {'日期': AC, '摘要': AL, '往来单位': AL, '经手人': AC, '收支项目': AL, '账户': AC, '备注': AL, '来源': AL}

    # 合计行（清单上方）
    sum_cols = [f for f in ('收入', '支出', '应收', '应付', '烘干重量', '购进数量', '销售数量') if f in C]
    tot = {f: f'=SUM({C[f]}{R0}:{C[f]}{R_END})' for f in sum_cols}
    tot['日期'] = '本期合计'
    tot['摘要'] = f'="共 "&{n}&" 条"&IF({n}>{SHOW},"（只显示了前 {SHOW} 条，合计只算显示的）","")'
    chk = [f'ROUND({C["收入"]}{R_TOT}-{S("本期收入")},2)=0', f'ROUND({C["支出"]}{R_TOT}-{S("本期支出")},2)=0',
           f'ROUND({C["应收"]}{R_TOT}-{S("本期应收增减")},2)=0', f'ROUND({C["应付"]}{R_TOT}-{S("本期应付增减")},2)=0']
    if has_stock:
        chk += [f'ROUND({C["购进数量"]}{R_TOT}-{S("本期购进")},3)=0', f'ROUND({C["销售数量"]}{R_TOT}-{S("本期销售")},3)=0']
    tot['来源'] = (f'=IF({n}>{SHOW},"⚠ 清单没列全，合计只算了显示的",IF(AND({",".join(chk)}),"✓ 合计＝顶部汇总",'
                 f'"⚠ 合计跟顶部汇总对不上"))')
    # 期初行
    opn = {'日期': f'=IF({S("起")}>0,{S("起")},"")', '摘要': '期初', '结余': f'={S("期初结余")}'}
    if has_stock:
        opn['存量'] = f'={S("期初存量")}'
    for row, vals, fl in ((R_TOT, tot, FILL_TOT), (R_OPEN, opn, FILL_SUB)):
        for f, c in C.items():
            put(ws, f'{c}{row}', vals.get(f), F_TXTB, fl, fmts.get(f, 'General'), aligns.get(f, AR))
    ws[f'{C["来源"]}{R_TOT}'].font = F_TXTB
    ws.conditional_formatting.add(f'{C["来源"]}{R_TOT}', FormulaRule(formula=[f'LEFT(${C["来源"]}${R_TOT},1)="⚠"'], font=F_RED))

    # 逐行
    for i in range(SHOW):
        r = R0 + i
        k, z = f'${KC}{r}', f'${ZC}{r}'
        ws[f'{KC}{r}'] = i + 1
        ws[f'{ZC}{r}'] = (f'=IF({k}>{n},0,MOD(IF({S("经手人")}="",SMALL({HK1},{S("起点")}+{k}),'
                          f'SMALL({HK2},{S("起点")}+{k})),10000))')
        ws[f'{KC}{r}'].font = ws[f'{ZC}{r}'].font = F_HELP
        zs, zw = z, f'{z}-{N_CASH}'

        def qnote(t, zz):
            """摘要后面加（数量KG×单价）：通用表有数量就加；购销表只给不是购进/销售的行加；烘干表不加（有专门的列）"""
            q, p = f'INDEX({t}_数量,{zz})', f'INDEX({t}_单价,{zz})'
            txt = f'"（"&{q}&"KG"&IF({p}=0,"","×"&{p})&"）"'
            if kind == '通用':
                return f'&IF({q}<>0,{txt},"")'
            if kind == '购销':
                s_ = f'INDEX({t}_库存,{zz})'
                return f'&IF(AND({s_}<>"入库",{s_}<>"出库",{q}<>0),{txt},"")'
            return ''

        def note(t, zz):
            v, bz = f'INDEX({t}_校验,{zz})', f'INDEX({t}_备注,{zz})'
            return f'{bz}&IF(OR(LEFT({v},1)="✗",LEFT({v},1)="⚠"),IF({bz}="","","；")&{v},"")'

        key = g('排序键', z)
        pe, a = S('板块条件'), S('起')
        upto = lambda t, fld: f'SUMIFS({t}_{fld},{t}_板块,{pe},{t}_日期,">="&{a},{t}_排序键,"<="&{key})'
        prev = lambda f: f'{C[f]}{r - 1}'
        v = {
            '日期': f'=IF({z}=0,"",{g("日期", z)})',
            '摘要': (f'=IF({z}=0,IF({k}={n}+1,IF({n}=0,"（没有流水）","共 "&{n}&" 条"),""),'
                   f'IF({z}<={N_CASH},INDEX(收_显示摘要,{zs}){qnote("收", zs)},'
                   f'INDEX(往_显示摘要,{zw}){qnote("往", zw)}&"（赊）"))'),
            '往来单位': f'=IF({z}=0,"",{g("往来单位", z)})',
            '收入': f'=IF({z}=0,"",IF({z}<={N_CASH},INDEX(收_板块收入,{zs}),""))',
            '支出': f'=IF({z}=0,"",IF({z}<={N_CASH},INDEX(收_板块支出,{zs}),""))',
            '结余': (f'=IF({z}=0,"",IF({S("经手人")}="",ROUND(N({prev("结余")})+N({C["收入"]}{r})-N({C["支出"]}{r}),2),'
                   f'ROUND({S("期初结余")}+{upto("收", "板块收入")}-{upto("收", "板块支出")},2)))'),
            '经手人': f'=IF({z}=0,"",{g("经手人", z)})',
            '收支项目': f'=IF({z}=0,"",{g("收支项目", z)})',
            '应收': f'=IF({z}=0,"",IF({z}<={N_CASH},-INDEX(收_冲应收,{zs}),INDEX(往_应收额,{zw})))',
            '应付': f'=IF({z}=0,"",IF({z}<={N_CASH},-INDEX(收_冲应付,{zs}),INDEX(往_应付额,{zw})))',
            '账户': f'=IF({z}=0,"",IF({z}<={N_CASH},INDEX(收_账户,{zs}),""))',
            '备注': f'=IF({z}=0,"",IF({z}<={N_CASH},{note("收", zs)},{note("往", zw)}))',
            '来源': (f'=IF({z}=0,"",IF({z}<={N_CASH},"收支登记第 "&INDEX(收_录入行,{zs})&" 行",'
                   f'"应收应付登记第 "&INDEX(往_录入行,{zw})&" 行"))'),
        }
        if kind == '烘干':
            v['烘干重量'] = f'=IF({z}=0,"",{g("数量", z)})'
            v['烘干单价'] = f'=IF({z}=0,"",IF(N({C["烘干重量"]}{r})=0,"",{g("单价", z)}))'
        if kind == '购销':
            v['购进数量'] = f'=IF({z}=0,"",{g("入库量", z)})'
            v['购进单价'] = f'=IF({z}=0,"",IF({g("库存", z)}="入库",{g("单价", z)},""))'
            v['销售数量'] = f'=IF({z}=0,"",{g("出库量", z)})'
            v['销售单价'] = f'=IF({z}=0,"",IF({g("库存", z)}="出库",{g("单价", z)},""))'
            v['存量'] = (f'=IF({z}=0,"",IF({S("经手人")}="",ROUND(N({prev("存量")})+N({C["购进数量"]}{r})-N({C["销售数量"]}{r}),3),'
                       f'ROUND({S("期初存量")}+{upto("收", "入库量")}+{upto("往", "入库量")}'
                       f'-{upto("收", "出库量")}-{upto("往", "出库量")},3)))')
        for f, c in C.items():
            q = ws[f'{c}{r}']
            q.value = v[f]
            q.font = F_TXT
            q.number_format = fmts.get(f, 'General')
            q.alignment = aligns.get(f, AR)

    rng = f'A{R0}:{last}{R_END}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${ZC}{R0}>{N_CASH}'], fill=fill(FILL_SHE), border=BD))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'${ZC}{R0}>0'], border=BD))
    ws.conditional_formatting.add(f'{C["摘要"]}{R0}:{C["摘要"]}{R_END}',
                                  FormulaRule(formula=[f'${ZC}{R0}=0'], font=Font(color='FF808080', italic=True)))
    ws.conditional_formatting.add(f'{C["备注"]}{R0}:{C["备注"]}{R_END}',
                                  FormulaRule(formula=[f'ISNUMBER(FIND("✗",{C["备注"]}{R0}))'], font=Font(color='FFC00000', bold=True)))

    hide(ws, KC, ZC, SC_COL, SC_LBL)
    ws.freeze_panes = f'C{R_TOT}'
    print_setup(ws, f'{HDR}:{HDR}', landscape=True)
    ws.print_area = f'A1:{last}{R_END}'


def build(wb, ctx=None):
    build_keys(wb[HK_SHEET])
    for seg in SEGMENTS:
        build_seg(wb[seg], seg, SEG_KIND[seg])
