# -*- coding: utf-8 -*-
"""查看表（绿）【查询】：照用户原表「查询」那一行的格式（第 2 行表头列不变），第 3 行选条件出结果，第 4 行补充金额和实际起止。
   下面：① 各板块一览 ② 按月 ③ 收支明细 ④ 应收应付明细。
   口径（设计.md §3）：
   - 粮食购进/销售＝期间数量 KG（收支登记＋应收应付登记），金额在第 4 行；
   - 收入/支出＝期间 收_板块收入/收_板块支出（不含内部转账）；
   - 粮食存量、结余、固定资产、应收、应付＝到截止日的余额（不按经手人分）；
   - 经手人只筛期间流水（购进、销售、收入、支出）和明细。
   起始日期空＝P_建账日；截止日期空＝P_截止。
   清单都是「活动行」：隐藏列先算出这一行显示第几条 / 合计 / 共 N 条，合计紧跟在最后一条下面。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from layout import *
from common import *

LAST = 'N'
GREEN_H = 'FF70AD47'
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'          # 零不显示（明细用）
KG = '#,##0.##;[Red]-#,##0.##;"-"'
KG_B = '#,##0.##;[Red]-#,##0.##;""'
PRICE_B = '#,##0.00##;[Red]-#,##0.00##;""'
MONTH_FMT = 'yyyy"年"m"月"'
F_GRAY_I = Font(name=YH, sz=10, italic=True, color='FF808080')
F_GRAY = Font(name=YH, sz=9, color='FF808080')


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


CF_TOT = cf_fill('FFFCE4D6')
CF_XFER = cf_fill('FFDDEBF7')
CF_PICK = cf_fill('FFFFF2CC')

# ── 标量（隐藏 AA 列，标签在 AB 列） ──
SEG, PER = '$AA$1', '$AA$2'                     # 选的板块、经手人（TRIM 后；空＝全部）
S0, S1 = '$B$4', '$C$4'                         # 实际起止
NM, NMS = '$AA$3', '$AA$4'                      # 起止跨几个月；显示几个月（≤24）
OTH, NSV, P_OTH, P_TOT = '$AA$5', '$AA$6', '$AA$7', '$AA$8'   # ①：「其他」要不要显示、显示几个板块、其他/合计排第几行
SEG_OK = '$AA$9'
XN, XIN, XOUT = '$AA$10', '$AA$11', '$AA$12'    # ③：内部转账条数、收、支
NN3, NN4 = '$AA$13', '$AA$14'                   # ③④：实际显示几条（≤显示上限）

# ── ① 隐藏计算块：第 i 个板块在 HR0+i-1 行；之后一行「其他」、一行「合计（全部）」 ──
HR0 = 5
HB = dict(名称='AD', 期初结余='AE', 收入='AF', 支出='AG', 期末结余='AH', 购量='AI', 购额='AJ', 销量='AK', 销额='AL',
          存量='AM', 固定资产='AN', 应收='AO', 应付='AP')
V1 = dict(名称='A', 期初结余='B', 收入='C', 支出='D', 期末结余='E', 购量='F', 购额='G', 销量='H', 销额='I',
          存量='J', 固定资产='K', 应收='L', 应付='M')
SEGCNT = 'AC'
I_OTH, I_TOT = N_SEG + 1, N_SEG + 2

IDX, KIND = 'P', 'Q'                            # 清单行：第几条（源表 n）、行类型（1 数据 2 合计 3… ）
KEY3, KEY4 = 'S', 'T'                           # ③④ 排序键
MF, MA, MB = 'U', 'V', 'W'                      # ②：月初、这个月实际起、实际止
SHOW_M, SHOW3, SHOW4 = 24, 800, 300

IN_S, OUT_S = '收_库存,"入库"', '收_库存,"出库"'
IN_W, OUT_W = '往_库存,"入库"', '往_库存,"出库"'


def S(rng, *c):
    return f'SUMIFS({rng},{",".join(x for x in c if x)})'


def R2(x):
    return f'ROUND({x},2)'


def flows(sS, sW, pS, pW, a, b):
    """期间流水：板块条件 sS/sW、经手人条件 pS/pW（None＝不筛）"""
    dS, dW = dr('收_日期', a, b), dr('往_日期', a, b)
    return dict(
        收入=S('收_板块收入', sS, pS, dS),
        支出=S('收_板块支出', sS, pS, dS),
        购量=f'{S("收_入库量", sS, pS, dS)}+{S("往_入库量", sW, pW, dW)}',
        购额=f'{S("收_支出", IN_S, sS, pS, dS)}+{S("往_金额", IN_W, "往_有效,1", sW, pW, dW)}',
        销量=f'{S("收_出库量", sS, pS, dS)}+{S("往_出库量", sW, pW, dW)}',
        销额=f'{S("收_收入", OUT_S, sS, pS, dS)}+{S("往_金额", OUT_W, "往_有效,1", sW, pW, dW)}',
    )


def bals(base, sS, sW, op, d):
    """到某天的余额：op＝"<="（到 d 为止）或 "<"（d 以前）；base(k)＝板块期初值"""
    tS, tW = f'收_日期,"{op}"&{d}', f'往_日期,"{op}"&{d}'
    return dict(
        结余=f'{base("期初结余")}+{S("收_板块收入", sS, tS)}-{S("收_板块支出", sS, tS)}',
        存量=(f'{base("期初存量")}+{S("收_入库量", sS, tS)}+{S("往_入库量", sW, tW)}'
            f'-{S("收_出库量", sS, tS)}-{S("往_出库量", sW, tW)}'),
        固定资产=f'{base("期初固定资产")}+{S("收_固定资产", sS, tS)}+{S("往_固定资产", sW, tW)}',
        应收=f'{S("往_应收额", sW, tW)}-{S("收_冲应收", sS, tS)}',
        应付=f'{S("往_应付额", sW, tW)}-{S("收_冲应付", sS, tS)}',
    )


def _sc(ws, cell, f, label):
    ws[cell.replace('$', '')] = f
    ws[cell.replace('$', '')].font = F_HELP
    lc = 'AB' + cell.replace('$', '')[2:]
    ws[lc] = label
    ws[lc].font = F_HELP


def _cell(ws, coord, v, fmt=None, align=AR, font=F_TXT, fill_=None, border=True):
    return put(ws, coord, v, font, fill_, fmt, align, border)


def _style_list(ws, r0, r1, cols, fmts, aligns):
    for r in range(r0, r1 + 1):
        for c in cols:
            x = ws[f'{c}{r}']
            x.font = F_TXT
            x.number_format = fmts.get(c, 'General')
            x.alignment = aligns.get(c, AC)


def build(wb, ctx):
    ctx = ctx or {}
    ws = wb[SH_QRY]
    nseg = sum(1 for s in ctx.get('segments', []) if str(s.get('名称', '') or '').strip())
    M1 = min(N_SEG, nseg + 2)

    widths(ws, {'A': 14, 'B': 13, 'C': 13, 'D': 20, 'E': 14, 'F': 14, 'G': 14, 'H': 13, 'I': 13, 'J': 14, 'K': 14,
                'L': 13, 'M': 13, 'N': 26})
    title(ws, '查  询（按日期、业务板块、经手人）', 'M', C_VIEW)
    home_link(ws, 'N1')

    # ═════ 第 2～4 行：用户原表的查询行 ═════
    heads = [('A', '查询'), ('B', '起始日期'), ('C', '截止日期'), ('D', '所属经营类别'), ('E', '粮食购进'), ('F', '粮食销售'),
             ('G', '粮食存量'), ('H', '收入'), ('I', '支出'), ('J', '结余'), ('K', '固定资产'), ('L', '应收'), ('M', '应付'),
             ('N', '经手人')]
    selector(ws, 'B2', '起始日期', 'B3', None, fmt=DATE)
    selector(ws, 'C2', '截止日期', 'C3', None, fmt=DATE)
    selector(ws, 'D2', '所属经营类别', 'D3', None, '=板块列表', prompt='从下拉选业务板块；空着＝全部板块')
    selector(ws, 'N2', '经手人', 'N3', None, '=经手人列表', prompt='从下拉选经手人；空着＝全部经手人')
    dv_date(ws, 'B3')
    dv_date(ws, 'C3')
    header(ws, 2, heads, C_VIEW, height=30)
    put(ws, 'A3', '本次查询 →', F_KPI_L, fill('FFD9E1F2'), align=AC)
    ws.row_dimensions[3].height = 28

    # 标量
    _sc(ws, SEG, '=TRIM(D3&"")', '选的板块')
    _sc(ws, PER, '=TRIM(N3&"")', '选的经手人')
    _sc(ws, NM, f'=IF({S0}>{S1},0,(YEAR({S1})-YEAR({S0}))*12+MONTH({S1})-MONTH({S0})+1)', '起止跨几个月')
    _sc(ws, NMS, f'=MIN({NM},{SHOW_M})', '按月显示几个月')
    _sc(ws, SEG_OK, f'=IF({SEG}="",1,IF(ISNUMBER(MATCH({SEG},板块_名称,0)),1,0))', '板块在基础资料里')

    segS, segW = seg_crit('收_板块', SEG), seg_crit('往_板块', SEG)
    perS, perW = seg_crit('收_经手人', PER), seg_crit('往_经手人', PER)
    base3 = lambda k: f'SUMIFS(板块_{k},{seg_crit("板块_名称", SEG)})'
    fl = flows(segS, segW, perS, perW, S0, S1)
    bl = bals(base3, segS, segW, '<=', S1)
    row3 = {'E': fl['购量'], 'F': fl['销量'], 'G': bl['存量'], 'H': fl['收入'], 'I': fl['支出'], 'J': bl['结余'],
            'K': bl['固定资产'], 'L': bl['应收'], 'M': bl['应付']}
    for c, f in row3.items():
        fm = KG if c in 'EFG' else MONEY
        _cell(ws, f'{c}3', f'={R2(f)}', fm, AR, F_KPI_V, FILL_AUTO)
    ws['E3'].number_format = ws['F3'].number_format = ws['G3'].number_format = '#,##0.##" KG";[Red]-#,##0.##" KG";"-"'

    # 第 4 行：实际用的、金额、说明
    put(ws, 'A4', '实际用的 / 金额', F_NOTE, fill('FFF2F2F2'), align=AC)
    _cell(ws, 'B4', '=IF(ISNUMBER(B3),INT(B3),P_建账日)', DATE, AC, F_AUTOB, FILL_AUTO)
    _cell(ws, 'C4', '=IF(ISNUMBER(C3),INT(C3),P_截止)', DATE, AC, F_AUTOB, FILL_AUTO)
    _cell(ws, 'D4', f'=IF({SEG}="","全部板块",{SEG}&IF({SEG_OK}=0,"（⚠ 不在基础资料里）",""))', None, AC, F_AUTOB, FILL_AUTO)
    _cell(ws, 'E4', f'={R2(fl["购额"])}', '"金额 "#,##0.00;[Red]"金额 "-#,##0.00;"金额 -"', AR, F_AUTOB, FILL_AUTO)
    _cell(ws, 'F4', f'={R2(fl["销额"])}', '"金额 "#,##0.00;[Red]"金额 "-#,##0.00;"金额 -"', AR, F_AUTOB, FILL_AUTO)
    ws.merge_cells('G4:M4')
    _cell(ws, 'G4', (f'="↑ 购进、销售、收入、支出＝"&TEXT({S0},"yyyy/m/d")&"～"&TEXT({S1},"yyyy/m/d")&" 这段时间的；'
                     f'存量、结余、固定资产、应收、应付＝到 "&TEXT({S1},"yyyy/m/d")&" 的余额"'
                     f'&IF({PER}="","","（不按经手人分）")'), None, AL, F_NOTE)
    for c in 'HIJKLM':
        ws[f'{c}4'].border = BD
    _cell(ws, 'N4', f'=IF({PER}="","全部经手人",{PER})', None, AC, F_AUTOB, FILL_AUTO)
    ws.row_dimensions[4].height = 22

    # 第 5 行：说明；第 6 行：提醒
    tip = ('💡 第 3 行黄格是查询条件：起始日期、截止日期（空着＝建账日～首页截止日）、所属经营类别（空＝全部板块）、经手人（空＝全部），'
           '选好自动出结果。粮食购进、粮食销售是这段时间的数量（KG，收支登记和应收应付登记都算），金额在第 4 行；收入、支出是这段时间的'
           '（不含账户之间倒钱）；粮食存量、结余、固定资产、应收、应付是到截止日期的余额。选了经手人：只筛购进、销售、收入、支出和下面的明细，'
           '余额类不按经手人分。下面：① 各板块一览 ② 按月 ③ 收支明细 ④ 应收应付明细（要改哪一笔，按最后一列行号去登记表改）。')
    ws.merge_cells(f'A5:{LAST}5')
    put(ws, 'A5', tip, F_TIP, FILL_TIP, align=ALW, border=False)
    ws.row_dimensions[5].height = 52
    ws.merge_cells(f'A6:{LAST}6')
    put(ws, 'A6', (f'=IF({S0}>{S1},"⚠ 起始日期晚于截止日期，请改日期　","")'
                   f'&IF({SEG_OK}=0,"⚠ 「"&{SEG}&"」不在【基础资料】的业务板块里　","")'
                   f'&IF(AND({PER}<>"",COUNTIF(收_经手人,{esc(PER)})+COUNTIF(往_经手人,{esc(PER)})=0),'
                   f'"⚠ 两张登记表里没有经手人是「"&{PER}&"」的","")'), F_RED, align=AL, border=False)

    # ═════ ① 各板块一览 ═════
    # 隐藏计算块
    for i in range(1, N_SEG + 1):
        r = HR0 + i - 1
        a = f'${HB["名称"]}{r}'
        e = esc(a)
        sS, sW = f'收_板块,{e}', f'往_板块,{e}'
        bI = lambda k, i=i: f'INDEX(板块_{k},{i})'
        f = flows(sS, sW, perS, perW, S0, S1)
        f.update(bals(bI, sS, sW, '<=', S1))
        f['期末结余'] = f.pop('结余')
        f['期初结余'] = bals(bI, sS, sW, '<', S0)['结余']
        ws[f'{HB["名称"]}{r}'] = f'=INDEX(板块_名称,{i})&""'
        for k, col in HB.items():
            if k != '名称':
                ws[f'{col}{r}'] = f'=IF({a}="",0,{R2(f[k])})'
    rt = HR0 + I_TOT - 1
    ro = HR0 + I_OTH - 1
    ball = lambda k: f'SUM(板块_{k})'
    f = flows(None, None, perS, perW, S0, S1)
    f.update(bals(ball, None, None, '<=', S1))
    f['期末结余'] = f.pop('结余')
    f['期初结余'] = bals(ball, None, None, '<', S0)['结余']
    ws[f'{HB["名称"]}{rt}'] = '合计'
    ws[f'{HB["名称"]}{ro}'] = '没选板块 / 板块名不对'
    for k, col in HB.items():
        if k == '名称':
            continue
        ws[f'{col}{rt}'] = f'={R2(f[k])}'
        ws[f'{col}{ro}'] = f'=ROUND({col}{rt}-SUM({col}{HR0}:{col}{HR0 + N_SEG - 1}),2)'
    for r in range(HR0, rt + 1):
        for col in HB.values():
            ws[f'{col}{r}'].font = F_HELP
    counter(ws, SEGCNT, HR0, N_SEG, lambda i: f'${HB["名称"]}{HR0 + i}<>""')
    nseg_cell = cnt(SEGCNT, HR0, N_SEG)
    _sc(ws, OTH, '=IF(OR(' + ','.join(f'ROUND({c}{ro},2)<>0' for k, c in HB.items() if k != '名称') + '),1,0)',
        '①「其他」要显示')
    _sc(ws, NSV, f'=MIN({nseg_cell},{M1})', '①显示几个板块')
    _sc(ws, P_OTH, f'=IF({OTH}=1,{NSV}+1,-1)', '①「其他」在第几行')
    _sc(ws, P_TOT, f'={NSV}+{OTH}+1', '①合计在第几行')

    r = 8
    section(ws, r, 'A', LAST, '① 各板块一览（起止、经手人同第 3 行；结余、存量、固定资产、应收应付是到截止日期的；选的那个板块黄底）', C_VIEW)
    r += 1
    header(ws, r, [('A', '业务板块'), ('B', '期初结余\n（起始日前一天）'), ('C', '收入'), ('D', '支出'), ('E', '期末结余\n（截止日）'),
                   ('F', '粮食购进\nKG'), ('G', '购进金额'), ('H', '粮食销售\nKG'), ('I', '销售金额'), ('J', '粮食存量\nKG'),
                   ('K', '固定资产'), ('L', '应收'), ('M', '应付'), ('N', '说明')], GREEN_H, height=36)
    r += 1
    a0 = r
    n1 = M1 + 2
    hr = lambda col: f'${col}${HR0}:${col}${HR0 + I_TOT - 1}'
    for k in range(1, n1 + 1):
        z = f'${IDX}{r}'
        ws[f'{IDX}{r}'] = (f'=IF({k}<={NSV},{kth(k, SEGCNT, HR0, N_SEG)},IF({k}={P_OTH},{I_OTH},'
                           f'IF({k}={P_TOT},{I_TOT},0)))')
        ws[f'{IDX}{r}'].font = F_HELP
        for key, c in V1.items():
            ws[f'{c}{r}'] = f'=IF({z}=0,"",INDEX({hr(HB[key])},{z}))'
        ws[f'N{r}'] = (f'=IF({z}=0,"",IF({z}<={N_SEG},IF(AND({SEG}<>"",A{r}={SEG}),"◀ 第 3 行选的就是这个",'
                       f'IF(E{r}<0,"⚠ 结余是负数","")),'
                       f'IF({z}={I_OTH},"⚠ 有收支没选板块或板块名不对（看【数据校验】）",'
                       f'IF({SEG}<>"","第 3 行只算「"&{SEG}&"」；选全部板块时第 3 行＝这一行",'
                       f'IF(AND(ROUND(C{r}-$H$3,2)=0,ROUND(D{r}-$I$3,2)=0,ROUND(E{r}-$J$3,2)=0,ROUND(F{r}-$E$3,2)=0,'
                       f'ROUND(G{r}-$E$4,2)=0,ROUND(H{r}-$F$3,2)=0,ROUND(I{r}-$F$4,2)=0,ROUND(J{r}-$G$3,2)=0,'
                       f'ROUND(K{r}-$K$3,2)=0,ROUND(L{r}-$L$3,2)=0,ROUND(M{r}-$M$3,2)=0),"✓ ＝第 3 行（全部板块）",'
                       f'"⚠ 跟第 3 行对不上")))&IF(AND({z}={I_TOT},{nseg_cell}>{M1}),'
                       f'"　⚠ 共 "&{nseg_cell}&" 个板块，只列了前 {M1} 个（合计是全部的）","")))')
        r += 1
    a1 = r - 1
    fm1 = {c: MONEY for c in 'BCDEGIKLM'}
    fm1.update({c: KG for c in 'FHJ'})
    _style_list(ws, a0, a1, list('ABCDEFGHIJKLMN'), fm1, {'A': AL, 'N': AL, **{c: AR for c in 'BCDEFGHIJKLM'}})
    rng1 = f'A{a0}:{LAST}{a1}'
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}={I_TOT}'], fill=CF_TOT, font=Font(bold=True),
                                                    border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}={I_OTH}'], font=Font(italic=True, color='FFC00000'),
                                                    border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'AND(${IDX}{a0}>0,{SEG}<>"",$A{a0}={SEG})'], fill=CF_PICK,
                                                    border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}>0'], border=BD))
    r += 1
    ws.merge_cells(f'A{r}:{LAST}{r}')
    put(ws, f'A{r}', (f'=IF({PER}="","期初结余＋收入－支出＝期末结余；账户之间倒钱（内部转账）不算收支。",'
                      f'"选了经手人「"&{PER}&"」：收入、支出、购进、销售只算他经手的；期初/期末结余、存量、固定资产、应收、应付是全部的，'
                      f'所以期初＋收－支不一定等于期末。")'), F_NOTE, align=AL, border=False)
    sec1 = (a0, a1)
    r += 2

    # ═════ ② 按月 ═════
    section(ws, r, 'A', LAST, '② 按月（板块、经手人同第 3 行；头尾两个月只算起止日期以内的；最多列 24 个月）', C_VIEW)
    r += 1
    header(ws, r, [('A', '月份'), ('B', '收入'), ('C', '支出'), ('D', '净额\n（收－支）'), ('E', '粮食购进\nKG'),
                   ('F', '粮食销售\nKG'), ('G', '月末粮食存量\nKG'), ('H', '说明')], GREEN_H, height=36)
    r += 1
    m0 = r
    for k in range(1, SHOW_M + 2):
        q = f'${KIND}{r}'
        ws[f'{MF}{r}'] = f'=DATE(YEAR({S0}),MONTH({S0})+{k - 1},1)'
        ws[f'{MA}{r}'] = f'=MAX({S0},{MF}{r})'
        ws[f'{MB}{r}'] = f'=MIN({S1},DATE(YEAR({S0}),MONTH({S0})+{k},0))'
        ws[f'{KIND}{r}'] = f'=IF({k}<={NMS},1,IF({k}={NMS}+1,2,0))'
        for c in (MF, MA, MB, KIND):
            ws[f'{c}{r}'].font = F_HELP
        a, b = f'${MA}{r}', f'${MB}{r}'
        fm = flows(segS, segW, perS, perW, a, b)
        stock = bals(base3, segS, segW, '<=', b)['存量']
        above = lambda c: f'SUM({c}${m0}:{c}{r - 1})' if r > m0 else '0'
        ws[f'A{r}'] = f'=IF({q}=1,{MF}{r},IF({q}=2,"合计",""))'
        ws[f'B{r}'] = f'=IF({q}=0,"",IF({q}=1,{R2(fm["收入"])},{above("B")}))'
        ws[f'C{r}'] = f'=IF({q}=0,"",IF({q}=1,{R2(fm["支出"])},{above("C")}))'
        ws[f'D{r}'] = f'=IF({q}=0,"",ROUND(B{r}-C{r},2))'
        ws[f'E{r}'] = f'=IF({q}=0,"",IF({q}=1,{R2(fm["购量"])},{above("E")}))'
        ws[f'F{r}'] = f'=IF({q}=0,"",IF({q}=1,{R2(fm["销量"])},{above("F")}))'
        ws[f'G{r}'] = f'=IF({q}=0,"",IF({q}=1,{R2(stock)},{f"G{r - 1}" if r > m0 else "0"}))'
        ws[f'H{r}'] = (f'=IF({q}=1,IF(OR({a}>{MF}{r},{b}<DATE(YEAR({MF}{r}),MONTH({MF}{r})+1,0)),'
                       f'"只算 "&TEXT({a},"m/d")&"～"&TEXT({b},"m/d"),""),'
                       f'IF({q}=2,IF({NM}=0,"⚠ 起始日期晚于截止日期",IF({NM}>{SHOW_M},"⚠ 起止跨了 "&{NM}&" 个月，只列了前 {SHOW_M} 个月（合计也只是这些月）",'
                       f'IF(AND(ROUND(B{r}-$H$3,2)=0,ROUND(C{r}-$I$3,2)=0,ROUND(E{r}-$E$3,2)=0,ROUND(F{r}-$F$3,2)=0,'
                       f'ROUND(G{r}-$G$3,2)=0),"✓ ＝第 3 行","⚠ 跟第 3 行对不上"))),""))')
        r += 1
    m1 = r - 1
    fm2 = {'A': MONTH_FMT, 'B': MONEY, 'C': MONEY, 'D': MONEY, 'E': KG, 'F': KG, 'G': KG}
    _style_list(ws, m0, m1, list('ABCDEFGH'), fm2, {'A': AC, 'H': AL, **{c: AR for c in 'BCDEFG'}})
    for rr in range(m0, m1 + 1):
        ws[f'H{rr}'].font = F_NOTE
    rng2 = f'A{m0}:G{m1}'
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'${KIND}{m0}=2'], fill=CF_TOT, font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'${KIND}{m0}=1'], border=BD))
    sec2 = (m0, m1)
    r += 1

    # ═════ ③ 收支明细 ═════
    section(ws, r, 'A', LAST, '③ 收支明细（收支登记里符合上面起止、板块、经手人的，按日期排；账户之间倒钱只在「全部板块」时列出，蓝底）', C_VIEW)
    r += 1
    note3 = r
    r += 1
    hdr3 = r
    header(ws, r, [('A', '日期'), ('B', '业务板块'), ('C', '收支项目'), ('D', '摘要'), ('E', '往来单位'), ('F', '数量\nKG'),
                   ('G', '单价'), ('H', '收入'), ('I', '支出'), ('J', '账户'), ('K', '经手人'), ('L', '收支登记\n行号')],
           GREEN_H, height=36)
    r += 1
    d0 = r

    def cond3(i):
        n = i + 1
        g = lambda k: f'INDEX(收_{k},{n})'
        return (f'AND({g("有效")}=1,{g("日期")}>={S0},{g("日期")}<={S1},'
                f'OR({SEG}="",AND({g("板块")}={SEG},{g("内部转账")}=0)),OR({PER}="",{g("经手人")}={PER}))')
    n3 = skey(ws, KEY3, d0, N_CASH, cond3, lambda i: f'INDEX(收_日期,{i + 1})')
    _sc(ws, NN3, f'=MIN({n3},{SHOW3})', '③显示几条')
    pS3 = dr('收_日期', S0, S1)
    _sc(ws, XN, f'=COUNTIFS(收_内部转账,1,{perS},{pS3})', '③内部转账条数')
    _sc(ws, XIN, f'=SUMIFS(收_收入,收_内部转账,1,{perS},{pS3})', '③内部转账收')
    _sc(ws, XOUT, f'=SUMIFS(收_支出,收_内部转账,1,{perS},{pS3})', '③内部转账支')
    tin = f'IF({SEG}="",SUMIFS(收_收入,{perS},{pS3}),SUMIFS(收_板块收入,收_板块,{esc(SEG)},{perS},{pS3}))'
    tout = f'IF({SEG}="",SUMIFS(收_支出,{perS},{pS3}),SUMIFS(收_板块支出,收_板块,{esc(SEG)},{perS},{pS3}))'
    for k in range(1, SHOW3 + 3):
        y, q = f'${IDX}{r}', f'${KIND}{r}'
        ws[f'{IDX}{r}'] = '=' + ksorted(k, KEY3, d0, N_CASH, n3) if k <= SHOW3 else 0
        ws[f'{KIND}{r}'] = f'=IF({y}>0,1,IF({k}={NN3}+1,2,IF({k}={NN3}+2,3,0)))'
        ws[f'{IDX}{r}'].font = ws[f'{KIND}{r}'].font = F_HELP
        g = lambda f_: f'INDEX(收_{f_},{y})'
        ws[f'A{r}'] = f'=IF({q}=1,{g("日期")},IF({q}=2,"合计",IF({q}=3,"共 "&{n3}&" 条","")))'
        ws[f'B{r}'] = f'=IF({q}=1,IF({g("内部转账")}=1,"（内部转账）",{g("板块")}),"")'
        ws[f'C{r}'] = f'=IF({q}=1,{g("收支项目")},"")'
        ws[f'D{r}'] = f'=IF({q}=1,{g("摘要")},"")'
        ws[f'E{r}'] = f'=IF({q}=1,{g("往来单位")},"")'
        ws[f'F{r}'] = f'=IF({q}=1,IF({g("数量")}=0,"",{g("数量")}),"")'
        ws[f'G{r}'] = f'=IF({q}=1,IF({g("单价")}=0,"",{g("单价")}),"")'
        ws[f'H{r}'] = f'=IF({q}=1,{g("收入")},IF({q}=2,ROUND({tin},2),""))'
        ws[f'I{r}'] = f'=IF({q}=1,{g("支出")},IF({q}=2,ROUND({tout},2),""))'
        ws[f'J{r}'] = f'=IF({q}=1,{g("账户")},"")'
        ws[f'K{r}'] = f'=IF({q}=1,{g("经手人")},"")'
        ws[f'L{r}'] = f'=IF({q}=1,{g("录入行")},"")'
        r += 1
    d1 = r - 1
    fm3 = {'A': DATE, 'F': KG_B, 'G': PRICE_B, 'H': MONEY_B, 'I': MONEY_B, 'L': '0'}
    _style_list(ws, d0, d1, list('ABCDEFGHIJKL'), fm3,
                {'A': AC, 'B': AL, 'C': AL, 'D': AL, 'E': AL, 'F': AR, 'G': AR, 'H': AR, 'I': AR, 'J': AC, 'K': AC, 'L': AC})
    rng3 = f'A{d0}:L{d1}'
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=2'], fill=CF_TOT, font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=3'], font=Font(color='FF808080')))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'AND(${KIND}{d0}=1,$B{d0}="（内部转账）")'], fill=CF_XFER,
                                                    font=Font(italic=True, color='FF595959'), border=BD))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=1'], border=BD))
    tot3 = f'INDEX($H${d0}:$H${d1},{NN3}+1)'
    tot3o = f'INDEX($I${d0}:$I${d1},{NN3}+1)'
    ws.merge_cells(f'A{note3}:{LAST}{note3}')
    put(ws, f'A{note3}', (f'=IF({n3}=0,"这段时间没有符合条件的收支",'
                          f'"共 "&{n3}&" 条，按日期排（同一天按登记顺序）"'
                          f'&IF({n3}>{SHOW3},"　⚠ 只列了前 {SHOW3} 条（合计是全部的），请缩短日期范围","")'
                          f'&IF(AND({SEG}="",{XN}>0),"；其中账户之间倒钱（内部转账）"&{XN}&" 条，不算收支，合计扣掉它们＝第 3 行收入、支出"'
                          f'&IF(AND(ROUND({tot3}-{XIN}-$H$3,2)=0,ROUND({tot3o}-{XOUT}-$I$3,2)=0)," ✓"," ⚠ 对不上"),'
                          f'IF(AND(ROUND({tot3}-$H$3,2)=0,ROUND({tot3o}-$I$3,2)=0),"；合计＝第 3 行收入、支出 ✓",'
                          f'"；⚠ 合计跟第 3 行收入、支出对不上")))'), F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'A{note3}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",A{note3}))'], font=Font(color='FFC00000', bold=True)))
    sec3 = (note3, hdr3, d0, d1)
    r += 1

    # ═════ ④ 应收应付明细 ═════
    section(ws, r, 'A', LAST, '④ 应收应付明细（应收应付登记里符合上面起止、板块、经手人的赊销、赊购，按日期排）', C_VIEW)
    r += 1
    note4 = r
    r += 1
    hdr4 = r
    header(ws, r, [('A', '日期'), ('B', '业务板块'), ('C', '往来单位'), ('D', '应收 / 应付'), ('E', '业务内容'), ('F', '摘要'),
                   ('G', '数量\nKG'), ('H', '单价'), ('I', '金额'), ('J', '约定收/付款\n日期'), ('K', '还没结的\n（到首页截止日）'),
                   ('L', '经手人'), ('M', '应收应付登记\n行号')], GREEN_H, height=36)
    r += 1
    w0 = r

    def cond4(i):
        n = i + 1
        g = lambda k: f'INDEX(往_{k},{n})'
        return (f'AND({g("有效")}=1,{g("日期")}>={S0},{g("日期")}<={S1},'
                f'OR({SEG}="",{g("板块")}={SEG}),OR({PER}="",{g("经手人")}={PER}))')
    n4 = skey(ws, KEY4, w0, N_WL, cond4, lambda i: f'INDEX(往_日期,{i + 1})')
    _sc(ws, NN4, f'=MIN({n4},{SHOW4})', '④显示几条')
    pW4 = dr('往_日期', S0, S1)
    for k in range(1, SHOW4 + 4):
        y, q = f'${IDX}{r}', f'${KIND}{r}'
        ws[f'{IDX}{r}'] = '=' + ksorted(k, KEY4, w0, N_WL, n4) if k <= SHOW4 else 0
        ws[f'{KIND}{r}'] = f'=IF({y}>0,1,IF({k}={NN4}+1,2,IF({k}={NN4}+2,3,IF({k}={NN4}+3,4,0))))'
        ws[f'{IDX}{r}'].font = ws[f'{KIND}{r}'].font = F_HELP
        g = lambda f_: f'INDEX(往_{f_},{y})'
        tsum = lambda rng, t: f'ROUND(SUMIFS({rng},往_类型,"{t}",往_有效,1,{segW},{perW},{pW4}),2)'
        ws[f'A{r}'] = f'=IF({q}=1,{g("日期")},IF({q}=2,"应收合计",IF({q}=3,"应付合计",IF({q}=4,"共 "&{n4}&" 条",""))))'
        ws[f'B{r}'] = f'=IF({q}=1,{g("板块")},"")'
        ws[f'C{r}'] = f'=IF({q}=1,{g("往来单位")},"")'
        ws[f'D{r}'] = f'=IF({q}=1,{g("类型")},"")'
        ws[f'E{r}'] = f'=IF({q}=1,{g("收支项目")},"")'
        ws[f'F{r}'] = f'=IF({q}=1,{g("摘要")},"")'
        ws[f'G{r}'] = f'=IF({q}=1,IF({g("数量")}=0,"",{g("数量")}),"")'
        ws[f'H{r}'] = f'=IF({q}=1,IF({g("单价")}=0,"",{g("单价")}),"")'
        ws[f'I{r}'] = (f'=IF({q}=1,{g("金额")},IF({q}=2,{tsum("往_金额", "应收")},'
                       f'IF({q}=3,{tsum("往_金额", "应付")},"")))')
        ws[f'J{r}'] = f'=IF({q}=1,{g("约定日期")},"")'
        ws[f'K{r}'] = (f'=IF({q}=1,{g("未结")},IF({q}=2,{tsum("往_未结", "应收")},'
                       f'IF({q}=3,{tsum("往_未结", "应付")},"")))')
        ws[f'L{r}'] = f'=IF({q}=1,{g("经手人")},"")'
        ws[f'M{r}'] = f'=IF({q}=1,{g("录入行")},"")'
        r += 1
    e1 = r - 1
    fm4 = {'A': DATE, 'G': KG_B, 'H': PRICE_B, 'I': MONEY_B, 'J': DATE, 'K': MONEY_B, 'M': '0'}
    _style_list(ws, w0, e1, list('ABCDEFGHIJKLM'), fm4,
                {'A': AC, 'B': AL, 'C': AL, 'D': AC, 'E': AL, 'F': AL, 'G': AR, 'H': AR, 'I': AR, 'J': AC, 'K': AR, 'L': AC,
                 'M': AC})
    rng4 = f'A{w0}:M{e1}'
    ws.conditional_formatting.add(rng4, FormulaRule(formula=[f'OR(${KIND}{w0}=2,${KIND}{w0}=3)'], fill=CF_TOT,
                                                    font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng4, FormulaRule(formula=[f'${KIND}{w0}=4'], font=Font(color='FF808080')))
    ws.conditional_formatting.add(rng4, FormulaRule(formula=[f'AND(${KIND}{w0}=1,$I{w0}<0)'], font=Font(color='FFC00000'),
                                                    border=BD))
    ws.conditional_formatting.add(rng4, FormulaRule(formula=[f'${KIND}{w0}=1'], border=BD))
    ws.merge_cells(f'A{note4}:{LAST}{note4}')
    put(ws, f'A{note4}', (f'=IF({n4}=0,"这段时间没有符合条件的赊销、赊购",'
                          f'"共 "&{n4}&" 条，按日期排；金额是负数的是折让、冲减（红字）；「还没结的」按先欠先还算到首页截止日 "'
                          f'&TEXT(P_截止,"yyyy/m/d")&IF({n4}>{SHOW4},"　⚠ 只列了前 {SHOW4} 条（合计是全部的），请缩短日期范围",""))'),
        F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'A{note4}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",A{note4}))'], font=Font(color='FFC00000', bold=True)))
    sec4 = (note4, hdr4, w0, e1)

    for i in range(CI('P'), CI('CD') + 1):
        ws.column_dimensions[CL(i)].hidden = True
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:4', landscape=True)
    ws.print_area = f'A1:{LAST}{e1}'
    ws._a070 = dict(sec1=sec1, sec2=sec2, sec3=sec3, sec4=sec4)
    return ws._a070
