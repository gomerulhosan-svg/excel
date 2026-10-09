# -*- coding: utf-8 -*-
"""查看表（绿）【资金台账】出纳的现金 / 银行日记账。
   黄格：账户（空＝全部账户）、起（空＝P_年初）、止（空＝P_截止）、板块（空＝全部）。
   ① 各账户：期初（起前一天，含账户期初余额）、本期收入、本期支出、期末、笔数；板块只影响收入、支出、笔数。
   ② 按板块：各板块本期收入、支出、净额（不含内部转账）；再一行账户之间倒钱（内部转账）。②按所选账户算。
   ③ 明细：符合条件的收支登记行按日期排，余额从期初逐笔滚动（选了板块时＝账户的实际余额，用 SUMIFS 按排序键算）。
   口径（设计.md §3）：账户余额（到 d）＝账户_期初余额＋Σ收_净额（收_账户＝它，日期≤d；含内部转账）。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *

LAST = 'K'
GREEN_H = 'FF70AD47'
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


CF_TOT = cf_fill('FFFCE4D6')
CF_XFER = cf_fill('FFDDEBF7')
CF_PICK = cf_fill('FFFFF2CC')

ACC, SEG = '$AA$1', '$AA$2'
ACC_OK, SEG_OK = '$AA$3', '$AA$4'
S0, S1 = '$D$4', '$F$4'
OTH1, NSV1, P_OTH1, P_TOT1 = '$AA$5', '$AA$6', '$AA$7', '$AA$8'
OTH2, NSV2, P_OTH2, P_TOT2 = '$AA$9', '$AA$10', '$AA$11', '$AA$12'
NN3 = '$AA$13'
LASTR = '$AA$14'                                 # 明细最后一行（「共 N 笔」）的行号：打印区域到这里

HR0 = 5
HA = dict(名称='AD', 期初='AE', 收入='AF', 支出='AG', 期末='AH', 笔数='AI')        # ① 每个账户（第 i 个在 HR0+i-1 行）
ACNT = 'AC'
A_OTH, A_TOT = N_ACC + 1, N_ACC + 2
HS = dict(名称='AK', 收入='AL', 支出='AM', 净额='AN')                               # ② 每个板块
SCNT = 'AJ'
S_OTH, S_TOT, S_XFER, S_ALL = N_SEG + 1, N_SEG + 2, N_SEG + 3, N_SEG + 4
IDX, KIND = 'P', 'Q'
KEY = 'BA'                                       # ③ 排序键：折成 500 行一列的方块（BA:BJ），都在明细行数以内
BLK = 500
SHOW = 1500


def S(rng, *c):
    return f'SUMIFS({rng},{",".join(x for x in c if x)})'


def R2(x):
    return f'ROUND({x},2)'


def key_block(ws, c0, r0, n, cond, datef):
    """排序键折成 BLK 行一列的方块：第 i 条（0 起）在 c0 右边第 i//BLK 列、第 r0+i%BLK 行；放「日期×10000＋n」，不满足放空。
       返回（个数格, 方块区域）；SMALL 对整块取第 k 小，MOD(…,10000) 还原 n。"""
    i0 = CI(c0)
    ncol = -(-n // BLK)
    for i in range(n):
        c = f'{CL(i0 + i // BLK)}{r0 + i % BLK}'
        ws[c] = f'=IF({cond(i)},{datef(i)}*10000+{i + 1},"")'
        ws[c].font = F_HELP
    blk = f'${c0}${r0}:${CL(i0 + ncol - 1)}${r0 + BLK - 1}'
    ws[f'{c0}{r0 - 1}'] = f'=COUNT({blk})'
    ws[f'{c0}{r0 - 1}'].font = F_HELP
    return f'${c0}${r0 - 1}', blk


def _sc(ws, cell, f, label):
    c = cell.replace('$', '')
    ws[c] = f
    ws[c].font = F_HELP
    ws['AB' + c[2:]] = label
    ws['AB' + c[2:]].font = F_HELP


def _cell(ws, coord, v, fmt=None, align=AR, font=F_TXT, fill_=None, border=True):
    return put(ws, coord, v, font, fill_, fmt, align, border)


def _style_list(ws, r0, r1, cols, fmts, aligns):
    for r in range(r0, r1 + 1):
        for c in cols:
            x = ws[f'{c}{r}']
            x.font = F_TXT
            x.number_format = fmts.get(c, 'General')
            x.alignment = aligns.get(c, AC)


def build(wb, ctx):
    ctx = ctx or {}
    ws = wb[SH_FUND]
    nacc = sum(1 for a in ctx.get('accounts', []) if str(a.get('名称', '') or '').strip())
    nseg = sum(1 for s in ctx.get('segments', []) if str(s.get('名称', '') or '').strip())
    MA = min(N_ACC, nacc + 2)
    MS = min(N_SEG, nseg + 2)

    widths(ws, {'A': 11, 'B': 13, 'C': 13, 'D': 13, 'E': 24, 'F': 14, 'G': 13, 'H': 13, 'I': 14, 'J': 9, 'K': 10})
    tip = ('💡 出纳的现金、银行日记账。黄格：账户（空＝全部账户合起来）、起止日期（空＝首页年初～截止日）、板块（空＝全部）。'
           '① 每个账户的期初（起始日前一天，含建账时的期初余额）、本期收入、支出、期末余额；② 按板块的收支（不含账户之间倒钱）；'
           '③ 逐笔明细，余额从期初逐笔往下滚，可以拿来跟银行对账单、现金盘点逐笔对。现金存银行、银行取现（内部转账）在账户里是真进真出，'
           '选全部账户时一出一进、余额不变。板块只筛本期收入、支出、笔数和明细（选了板块就不算内部转账），不影响期初、期末（选了板块时明细的余额是账户的实际余额）。'
           '要改哪一笔，按最后一列行号去【收支登记】改。')
    title(ws, '资金台账（现金 / 银行日记账）', LAST, C_VIEW, tip)

    # ── 选择格 ──
    selector(ws, 'A3', '账户', 'B3', None, '=账户列表', prompt='从下拉选账户；空着＝全部账户合起来')
    selector(ws, 'C3', '起', 'D3', None, fmt=DATE)
    selector(ws, 'E3', '止', 'F3', None, fmt=DATE)
    selector(ws, 'G3', '板块', 'H3', None, '=板块列表', prompt='从下拉选业务板块；空着＝全部（含账户之间倒钱）')
    dv_date(ws, 'D3')
    dv_date(ws, 'F3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 26
    _sc(ws, ACC, '=TRIM(B3&"")', '选的账户')
    _sc(ws, SEG, '=TRIM(H3&"")', '选的板块')
    _sc(ws, ACC_OK, f'=IF({ACC}="",1,IF(ISNUMBER(MATCH({esc(ACC)},账户_名称,0)),1,0))', '账户在基础资料里')
    _sc(ws, SEG_OK, f'=IF({SEG}="",1,IF(ISNUMBER(MATCH({esc(SEG)},板块_名称,0)),1,0))', '板块在基础资料里')
    put(ws, 'A4', '实际用的', F_NOTE, align=AC)
    _cell(ws, 'B4', f'=IF({ACC}="","全部账户",{ACC})', None, AC, F_AUTOB, FILL_AUTO)
    put(ws, 'C4', '空＝首页年初→', F_NOTE, align=AR)
    _cell(ws, 'D4', '=IF(ISNUMBER(D3),INT(D3),P_年初)', DATE, AC, F_AUTOB, FILL_AUTO)
    put(ws, 'E4', '空＝首页截止日→', F_NOTE, align=AR)
    _cell(ws, 'F4', '=IF(ISNUMBER(F3),INT(F3),P_截止)', DATE, AC, F_AUTOB, FILL_AUTO)
    put(ws, 'G4', '', F_NOTE, align=AR)
    _cell(ws, 'H4', f'=IF({SEG}="","全部板块",{SEG})', None, AC, F_AUTOB, FILL_AUTO)
    ws.merge_cells(f'I4:{LAST}4')
    put(ws, 'I4', (f'=IF({S0}>{S1},"⚠ 起晚于止，请改日期　","")&IF({ACC_OK}=0,"⚠ 「"&{ACC}&"」不在【基础资料】的账户里　","")'
                   f'&IF({SEG_OK}=0,"⚠ 「"&{SEG}&"」不在【基础资料】的业务板块里","")'), F_RED, align=AL, border=False)

    accS, accN = seg_crit('收_账户', ACC), seg_crit('账户_名称', ACC)
    # 选了板块：只算这个板块的、而且不是内部转账的（内部转账就算填了板块也不算板块收支）；没选＝全部（含内部转账）
    segS = seg_crit('收_板块', SEG) + f',收_内部转账,IF({SEG}="","<>@@全部@@",0)'
    pS = dr('收_日期', S0, S1)

    # ── 第 5～6 行：所选账户本期汇总（关键结果格 C6:G6） ──
    ws.merge_cells('A5:B6')
    put(ws, 'A5', '="所选："&B4&IF(' + SEG + '="","","（"&' + SEG + '&"）")', F_KPI_L, fill('FFD9E1F2'), align=ACW)
    ws['B5'].border = ws['A6'].border = ws['B6'].border = BD
    for c, t in zip('CDEFG', ('期初余额\n（起始日前一天）', '本期收入', '本期支出', '期末余额', '笔数')):
        put(ws, f'{c}5', t, F_KPI_L, fill('FFD9E1F2'), align=ACW)
    lt0, le1 = f'收_日期,"<"&{S0}', f'收_日期,"<="&{S1}'
    _cell(ws, 'C6', f'={R2(S("账户_期初余额", accN) + "+" + S("收_净额", accS, lt0))}', MONEY, AR, F_KPI_V, FILL_AUTO)
    _cell(ws, 'D6', f'={R2(S("收_收入", accS, segS, pS))}', MONEY, AR, F_KPI_V, FILL_AUTO)
    _cell(ws, 'E6', f'={R2(S("收_支出", accS, segS, pS))}', MONEY, AR, F_KPI_V, FILL_AUTO)
    _cell(ws, 'F6', f'={R2(S("账户_期初余额", accN) + "+" + S("收_净额", accS, le1))}', MONEY, AR, F_KPI_V, FILL_AUTO)
    _cell(ws, 'G6', f'=COUNTIFS(收_有效,1,{accS},{segS},{pS})', INT, AC, F_KPI_V, FILL_AUTO)
    ws.merge_cells(f'H5:{LAST}6')
    put(ws, 'H5', (f'=IF({SEG}<>"","选了板块：收入、支出、笔数只算「"&{SEG}&"」；期初、期末是账户全部的",'
                   f'IF(ROUND(C6+D6-E6-F6,2)=0,"✓ 期初＋本期收入－本期支出＝期末","⚠ 期初＋收入－支出≠期末"))'),
        F_NOTE, align=ALW, border=False)
    ws.row_dimensions[5].height = 30
    ws.row_dimensions[6].height = 24

    # ═════ ① 各账户 ═════
    for i in range(1, N_ACC + 1):
        r = HR0 + i - 1
        a = f'${HA["名称"]}{r}'
        e = esc(a)
        f = {
            '名称': f'=INDEX(账户_名称,{i})&""',
            '期初': f'=IF({a}="",0,{R2(f"INDEX(账户_期初余额,{i})+" + S("收_净额", f"收_账户,{e}", lt0))})',
            '收入': f'=IF({a}="",0,{R2(S("收_收入", f"收_账户,{e}", segS, pS))})',
            '支出': f'=IF({a}="",0,{R2(S("收_支出", f"收_账户,{e}", segS, pS))})',
            '期末': f'=IF({a}="",0,{R2(f"INDEX(账户_期初余额,{i})+" + S("收_净额", f"收_账户,{e}", le1))})',
            '笔数': f'=IF({a}="",0,COUNTIFS(收_账户,{e},收_有效,1,{segS},{pS}))',
        }
        for k, v in f.items():
            ws[f'{HA[k]}{r}'] = v
            ws[f'{HA[k]}{r}'].font = F_HELP
    ro, rt = HR0 + A_OTH - 1, HR0 + A_TOT - 1
    ws[f'{HA["名称"]}{ro}'] = '账户名不对'
    ws[f'{HA["名称"]}{rt}'] = '合计'
    tot = {
        '期初': R2('SUM(账户_期初余额)+' + S('收_净额', f'收_日期,"<"&{S0}')),
        '收入': R2(S('收_收入', segS, pS)),
        '支出': R2(S('收_支出', segS, pS)),
        '期末': R2('SUM(账户_期初余额)+' + S('收_净额', f'收_日期,"<="&{S1}')),
        '笔数': f'COUNTIFS(收_有效,1,{segS},{pS})',
    }
    for k, v in tot.items():
        c = HA[k]
        ws[f'{c}{rt}'] = '=' + v
        ws[f'{c}{ro}'] = f'=ROUND({c}{rt}-SUM({c}{HR0}:{c}{HR0 + N_ACC - 1}),2)'
        ws[f'{c}{rt}'].font = ws[f'{c}{ro}'].font = F_HELP
    counter(ws, ACNT, HR0, N_ACC, lambda i: f'${HA["名称"]}{HR0 + i}<>""')
    nacc_cell = cnt(ACNT, HR0, N_ACC)
    _sc(ws, OTH1, '=IF(OR(' + ','.join(f'ROUND({HA[k]}{ro},2)<>0' for k in tot) + '),1,0)', '①「账户名不对」要显示')
    _sc(ws, NSV1, f'=MIN({nacc_cell},{MA})', '①显示几个账户')
    _sc(ws, P_OTH1, f'=IF({OTH1}=1,{NSV1}+1,-1)', '①「账户名不对」在第几行')
    _sc(ws, P_TOT1, f'={NSV1}+{OTH1}+1', '①合计在第几行')

    r = 8
    section(ws, r, 'A', LAST, '① 各账户（期初＝起始日前一天的余额；收入、支出含账户之间倒钱；选的账户黄底）', C_VIEW)
    r += 1
    header(ws, r, [('A', '账户'), ('C', '期初余额\n（起始日前一天）'), ('D', '本期收入'), ('E', '本期支出'), ('F', '期末余额'),
                   ('G', '笔数'), ('H', '说明')], GREEN_H, height=36)
    ws.merge_cells(f'A{r}:B{r}')
    ws.merge_cells(f'H{r}:{LAST}{r}')
    r += 1
    a0 = r
    hr = lambda col, n: f'${col}${HR0}:${col}${HR0 + n - 1}'
    for k in range(1, MA + 3):
        z = f'${IDX}{r}'
        ws[f'{IDX}{r}'] = (f'=IF({k}<={NSV1},{kth(k, ACNT, HR0, N_ACC)},IF({k}={P_OTH1},{A_OTH},'
                           f'IF({k}={P_TOT1},{A_TOT},0)))')
        ws[f'{IDX}{r}'].font = F_HELP
        ws.merge_cells(f'A{r}:B{r}')
        for c, key in (('A', '名称'), ('C', '期初'), ('D', '收入'), ('E', '支出'), ('F', '期末'), ('G', '笔数')):
            ws[f'{c}{r}'] = f'=IF({z}=0,"",INDEX({hr(HA[key], A_TOT)},{z}))'
        ws[f'H{r}'] = (f'=IF({z}=0,"",IF({z}<={N_ACC},IF(AND({ACC}<>"",A{r}={ACC}),"◀ 选的就是这个账户　","")'
                       f'&IF(F{r}<0,"⚠ 余额是负数：漏记收入或记错账户？",""),'
                       f'IF({z}={A_OTH},"⚠ 收支登记里有账户不在【基础资料】③里（看【数据校验】）",'
                       f'IF({SEG}<>"","板块只影响本期收入、支出、笔数；期初、期末是账户全部的",'
                       f'IF(ROUND(C{r}+D{r}-E{r}-F{r},2)=0,"✓ 期初＋收入－支出＝期末","⚠ 期初＋收入－支出≠期末"))'
                       f'&IF({nacc_cell}>{MA},"　⚠ 共 "&{nacc_cell}&" 个账户，只列了前 {MA} 个（合计是全部的）",""))))')
        r += 1
    a1 = r - 1
    _style_list(ws, a0, a1, list('ABCDEFGH'), {'C': MONEY, 'D': MONEY, 'E': MONEY, 'F': MONEY, 'G': INT},
                {'A': AL, 'B': AL, 'C': AR, 'D': AR, 'E': AR, 'F': AR, 'G': AC, 'H': AL})
    for rr in range(a0, a1 + 1):
        ws[f'H{rr}'].font = F_NOTE
    rng1 = f'A{a0}:G{a1}'
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}={A_TOT}'], fill=CF_TOT, font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}={A_OTH}'], font=Font(italic=True, color='FFC00000'),
                                                    border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'AND(${IDX}{a0}>0,{ACC}<>"",$A{a0}={ACC})'], fill=CF_PICK, border=BD))
    ws.conditional_formatting.add(rng1, FormulaRule(formula=[f'${IDX}{a0}>0'], border=BD))
    ws.conditional_formatting.add(f'H{a0}:H{a1}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",H{a0}))'],
                                                             font=Font(color='FFC00000', bold=True)))
    sec1 = (a0, a1)
    r += 1

    # ═════ ② 按板块 ═════
    for i in range(1, N_SEG + 1):
        rr = HR0 + i - 1
        a = f'${HS["名称"]}{rr}'
        e = esc(a)
        ws[f'{HS["名称"]}{rr}'] = f'=INDEX(板块_名称,{i})&""'
        ws[f'{HS["收入"]}{rr}'] = f'=IF({a}="",0,{R2(S("收_板块收入", f"收_板块,{e}", accS, pS))})'
        ws[f'{HS["支出"]}{rr}'] = f'=IF({a}="",0,{R2(S("收_板块支出", f"收_板块,{e}", accS, pS))})'
        ws[f'{HS["净额"]}{rr}'] = f'=ROUND({HS["收入"]}{rr}-{HS["支出"]}{rr},2)'
        for c in HS.values():
            ws[f'{c}{rr}'].font = F_HELP
    row_of = lambda idx: HR0 + idx - 1
    ro2, rt2, rx2, ra2 = row_of(S_OTH), row_of(S_TOT), row_of(S_XFER), row_of(S_ALL)
    ws[f'{HS["名称"]}{ro2}'] = '其他（没选板块 / 板块名不对）'
    ws[f'{HS["名称"]}{rt2}'] = '合计（不含内部转账）'
    ws[f'{HS["名称"]}{rx2}'] = '账户之间倒钱（内部转账）'
    ws[f'{HS["名称"]}{ra2}'] = '全部收支（含倒钱）'
    I, O, NE = HS['收入'], HS['支出'], HS['净额']
    ws[f'{I}{rt2}'] = f'={R2(S("收_板块收入", accS, pS))}'
    ws[f'{O}{rt2}'] = f'={R2(S("收_板块支出", accS, pS))}'
    ws[f'{I}{rx2}'] = f'={R2(S("收_收入", "收_内部转账,1", accS, pS))}'
    ws[f'{O}{rx2}'] = f'={R2(S("收_支出", "收_内部转账,1", accS, pS))}'
    for c in (I, O):
        ws[f'{c}{ro2}'] = f'=ROUND({c}{rt2}-SUM({c}{HR0}:{c}{HR0 + N_SEG - 1}),2)'
        ws[f'{c}{ra2}'] = f'=ROUND({c}{rt2}+{c}{rx2},2)'
    for rr in (ro2, rt2, rx2, ra2):
        ws[f'{NE}{rr}'] = f'=ROUND({I}{rr}-{O}{rr},2)'
        for c in HS.values():
            ws[f'{c}{rr}'].font = F_HELP
    counter(ws, SCNT, HR0, N_SEG, lambda i: f'${HS["名称"]}{HR0 + i}<>""')
    nseg_cell = cnt(SCNT, HR0, N_SEG)
    _sc(ws, OTH2, f'=IF(OR(ROUND({I}{ro2},2)<>0,ROUND({O}{ro2},2)<>0),1,0)', '②「其他」要显示')
    _sc(ws, NSV2, f'=MIN({nseg_cell},{MS})', '②显示几个板块')
    _sc(ws, P_OTH2, f'=IF({OTH2}=1,{NSV2}+1,-1)', '②「其他」在第几行')
    _sc(ws, P_TOT2, f'={NSV2}+{OTH2}+1', '②合计在第几行（倒钱、全部紧跟）')

    section(ws, r, 'A', LAST, '② 按板块（所选账户、起止；不含账户之间倒钱；下面再列倒钱，加起来＝①所选账户的本期收入、支出）', C_VIEW)
    r += 1
    header(ws, r, [('A', '业务板块'), ('C', '本期收入'), ('D', '本期支出'), ('E', '净额（收－支）'), ('F', '说明')], GREEN_H, height=30)
    ws.merge_cells(f'A{r}:B{r}')
    ws.merge_cells(f'F{r}:{LAST}{r}')
    r += 1
    b0 = r
    for k in range(1, MS + 5):
        z = f'${IDX}{r}'
        ws[f'{IDX}{r}'] = (f'=IF({k}<={NSV2},{kth(k, SCNT, HR0, N_SEG)},IF({k}={P_OTH2},{S_OTH},IF({k}={P_TOT2},{S_TOT},'
                           f'IF({k}={P_TOT2}+1,{S_XFER},IF({k}={P_TOT2}+2,{S_ALL},0)))))')
        ws[f'{IDX}{r}'].font = F_HELP
        ws.merge_cells(f'A{r}:B{r}')
        for c, key in (('A', '名称'), ('C', '收入'), ('D', '支出'), ('E', '净额')):
            ws[f'{c}{r}'] = f'=IF({z}=0,"",INDEX({hr(HS[key], S_ALL)},{z}))'
        ws[f'F{r}'] = (f'=IF({z}=0,"",IF({z}<={N_SEG},IF(AND({SEG}<>"",A{r}={SEG}),"◀ 选的就是这个板块",""),'
                       f'IF({z}={S_OTH},"⚠ 有收支没选板块或板块名不对（看【数据校验】）",'
                       f'IF({z}={S_XFER},"现金存银行、银行取现：一出一进，收、支各算一次，不是真的收支",'
                       f'IF({z}={S_ALL},IF({SEG}<>"","第 6 行只算「"&{SEG}&"」板块（＝上面「"&{SEG}&"」那一行）",'
                       f'IF(AND(ROUND(C{r}-$D$6,2)=0,ROUND(D{r}-$E$6,2)=0),"✓ ＝第 6 行所选账户的本期收入、支出",'
                       f'"⚠ 跟第 6 行本期收入、支出对不上")),'
                       f'IF({nseg_cell}>{MS},"⚠ 共 "&{nseg_cell}&" 个板块，只列了前 {MS} 个（合计是全部的）",""))))))')
        r += 1
    b1 = r - 1
    _style_list(ws, b0, b1, list('ABCDEF'), {'C': MONEY, 'D': MONEY, 'E': MONEY},
                {'A': AL, 'B': AL, 'C': AR, 'D': AR, 'E': AR, 'F': AL})
    for rr in range(b0, b1 + 1):
        ws[f'F{rr}'].font = F_NOTE
    rng2 = f'A{b0}:E{b1}'
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'OR(${IDX}{b0}={S_TOT},${IDX}{b0}={S_ALL})'], fill=CF_TOT,
                                                    font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'${IDX}{b0}={S_XFER}'], fill=CF_XFER,
                                                    font=Font(italic=True, color='FF595959'), border=BD))
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'${IDX}{b0}={S_OTH}'], font=Font(italic=True, color='FFC00000'),
                                                    border=BD))
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'AND(${IDX}{b0}>0,${IDX}{b0}<={N_SEG},{SEG}<>"",$A{b0}={SEG})'],
                                                    fill=CF_PICK, border=BD))
    ws.conditional_formatting.add(rng2, FormulaRule(formula=[f'${IDX}{b0}>0'], border=BD))
    ws.conditional_formatting.add(f'F{b0}:F{b1}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",F{b0}))'],
                                                             font=Font(color='FFC00000', bold=True)))
    sec2 = (b0, b1)
    r += 1

    # ═════ ③ 明细 ═════
    section(ws, r, 'A', LAST, '③ 逐笔明细（按日期排，同一天按登记顺序；余额从期初逐笔滚动；账户之间倒钱蓝底）', C_VIEW)
    r += 1
    note3 = r
    r += 1
    hdr3 = r
    header(ws, r, [('A', '日期'), ('B', '账户'), ('C', '业务板块'), ('D', '收支项目'), ('E', '摘要'), ('F', '往来单位'),
                   ('G', '收入'), ('H', '支出'), ('I', '余额'), ('J', '经手人'), ('K', '收支登记\n行号')], GREEN_H, height=36)
    r += 1
    d0 = r

    def cond(i):
        n = i + 1
        g = lambda k: f'INDEX(收_{k},{n})'
        return (f'AND({g("有效")}=1,{g("日期")}>={S0},{g("日期")}<={S1},'
                f'OR({ACC}="",{g("账户")}={ACC}),OR({SEG}="",AND({g("板块")}={SEG},{g("内部转账")}=0)))')
    n3, blk = key_block(ws, KEY, d0, N_CASH, cond, lambda i: f'INDEX(收_日期,{i + 1})')
    _sc(ws, NN3, f'=MIN({n3},{SHOW})', '③显示几条')
    _sc(ws, LASTR, f'={d0}+{NN3}+1', '明细最后一行（打印到这里）')
    base_acc = S('账户_期初余额', accN)
    for k in range(1, SHOW + 3):
        y, q = f'${IDX}{r}', f'${KIND}{r}'
        ws[f'{IDX}{r}'] = f'=IF({k}>{n3},0,MOD(SMALL({blk},{k}),10000))' if k <= SHOW else 0
        ws[f'{KIND}{r}'] = f'=IF({y}>0,1,IF({k}={NN3}+1,2,IF({k}={NN3}+2,3,0)))'
        ws[f'{IDX}{r}'].font = ws[f'{KIND}{r}'].font = F_HELP
        g = lambda f_: f'INDEX(收_{f_},{y})'
        prev = '$C$6' if k == 1 else f'I{r - 1}'
        ws[f'A{r}'] = f'=IF({q}=1,{g("日期")},IF({q}=2,"合计",IF({q}=3,"共 "&{n3}&" 笔","")))'
        ws[f'B{r}'] = f'=IF({q}=1,{g("账户")},"")'
        ws[f'C{r}'] = f'=IF({q}=1,IF({g("内部转账")}=1,"（内部转账）",{g("板块")}),"")'
        ws[f'D{r}'] = f'=IF({q}=1,{g("收支项目")},"")'
        ws[f'E{r}'] = f'=IF({q}=1,{g("摘要")},"")'
        ws[f'F{r}'] = f'=IF({q}=1,{g("往来单位")},"")'
        ws[f'G{r}'] = f'=IF({q}=1,{g("收入")},IF({q}=2,$D$6,""))'
        ws[f'H{r}'] = f'=IF({q}=1,{g("支出")},IF({q}=2,$E$6,""))'
        ws[f'I{r}'] = (f'=IF({q}=1,IF({SEG}="",ROUND({prev}+G{r}-H{r},2),'
                       f'ROUND({base_acc}+SUMIFS(收_净额,{accS},收_排序键,"<="&{g("排序键")}),2)),'
                       f'IF({q}=2,$F$6,""))')
        ws[f'J{r}'] = f'=IF({q}=1,{g("经手人")},"")'
        ws[f'K{r}'] = f'=IF({q}=1,{g("录入行")},"")'
        r += 1
    d1 = r - 1
    _style_list(ws, d0, d1, list('ABCDEFGHIJK'), {'A': DATE, 'G': MONEY_B, 'H': MONEY_B, 'I': MONEY, 'K': '0'},
                {'A': AC, 'B': AC, 'C': AL, 'D': AL, 'E': AL, 'F': AL, 'G': AR, 'H': AR, 'I': AR, 'J': AC, 'K': AC})
    rng3 = f'A{d0}:K{d1}'
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=2'], fill=CF_TOT, font=Font(bold=True), border=BD))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=3'], font=Font(color='FF808080')))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'AND(${KIND}{d0}=1,$C{d0}="（内部转账）")'], fill=CF_XFER,
                                                    border=BD))
    ws.conditional_formatting.add(rng3, FormulaRule(formula=[f'${KIND}{d0}=1'], border=BD))
    last_bal = f'INDEX($I${d0}:$I${d1},{NN3})'
    ws.merge_cells(f'A{note3}:{LAST}{note3}')
    put(ws, f'A{note3}', (f'=IF({n3}=0,IF({SEG}<>"","这段时间这个板块没有流水（上面的期初、期末是账户全部的）","这段时间没有流水（余额不变：期初＝期末）"),'
                          f'"共 "&{n3}&" 笔，按日期排（同一天按登记顺序）"'
                          f'&IF({n3}>{SHOW},"　⚠ 只显示前 {SHOW} 笔（余额也只滚到第 {SHOW} 笔），请缩短日期范围","")'
                          f'&IF({SEG}<>"","；选了板块：余额是账户的实际余额（含别的板块的收支）",'
                          f'IF(ROUND({last_bal}-$F$6,2)=0,"；最后一笔的余额＝期末余额 ✓","；⚠ 最后一笔的余额跟期末余额对不上")))'),
        F_NOTE, align=AL, border=False)
    ws.conditional_formatting.add(f'A{note3}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",A{note3}))'],
                                                          font=Font(color='FFC00000', bold=True)))
    sec3 = (note3, hdr3, d0, d1)

    for i in range(CI('P'), CI('CD') + 1):
        ws.column_dimensions[CL(i)].hidden = True
    ws.freeze_panes = 'A5'
    print_setup(ws, f'{hdr3}:{hdr3}', landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q_ = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName('Print_Area', attr_text=f'{q_}!$A$1:INDEX({q_}!${LAST}$1:${LAST}${d1},{q_}!{LASTR})')
    ws._a070 = dict(sec1=sec1, sec2=sec2, sec3=sec3)
    return ws._a070
