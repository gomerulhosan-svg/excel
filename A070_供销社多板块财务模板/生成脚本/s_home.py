# -*- coding: utf-8 -*-
"""【首页】：截止日期（全书默认）、关键数、各板块一览、提醒、各表入口、常见情况怎么记。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *
import s_check

NAV = [
    ('平时录（蓝）', C_IN, [(SH_CASH, '登记汇总表：所有板块的每一笔钱'), (SH_WL, '赊销、赊购（钱没当场收/付）')]),
    ('查看（绿，自动）', C_VIEW, [(SH_QRY, '按日期、板块、经手人查'), (SH_FUND, '现金、银行日记账')] +
     [(s, f'{s}的收支台账（自动拆分）') for s in SEGMENTS]),
    ('往来（橙，自动）', C_WL, [(SH_FOLLOW, '谁欠我们、我们欠谁、逾期没有'), (SH_STMT, '选单位自动出对账单，可打印')]),
    ('基础（灰）', C_MASTER, [(SH_BASE, '单位名称、板块、账户、收支项目、经手人'), (SH_UNIT, '客户、供应商（联系人、电话、账期）')]),
    ('校验', C_CHK, [(SH_CHK, '哪里录错了')]),
]

HOWTO = [
    ('收钱、付钱（当场结清的）', '【收支登记】一行：日期、业务板块、收支项目、摘要、收入或支出；卖粮买粮、烘干按斤算的填数量和单价，金额＝数量×单价'),
    ('卖出去没收到钱（赊销）', '【应收应付登记】记一行「应收」：往来单位、业务内容、数量单价或金额，可以填约定收款日期'),
    ('买进来没付钱（赊购）', '【应收应付登记】记一行「应付」'),
    ('收回以前的欠款', '【收支登记】收支项目选「收回欠款」，往来单位选欠钱的那家，业务板块选当初赊销的板块'),
    ('付掉以前欠人家的钱', '【收支登记】收支项目选「支付欠款」，往来单位选那家'),
    ('折让、少收、抹零', '【应收应付登记】同一个单位记一行负数'),
    ('现金存银行、银行取现', '【收支登记】记两行「内部转账」：转出的账户记支出，转入的账户记收入，金额一样；不算收支'),
    ('板块之间调钱', '【收支登记】记两行「板块调拨」：调出的板块记支出，调入的板块记收入'),
    ('买设备、房屋', '收支项目选「购置固定资产」（赊购的在应收应付登记记应付，业务内容也选它）；查询里「固定资产」就有了'),
    ('跟对方对账', '【对账单】选往来单位和日期，打印出来双方签字'),
]


def build(wb, ctx=None):
    ws = wb[SH_HOME]
    widths(ws, {'A': 2, 'B': 16, 'C': 15, 'D': 3, 'E': 16, 'F': 15, 'G': 3, 'H': 16, 'I': 15, 'J': 3, 'K': 16, 'L': 15})
    title(ws, '供 销 社 多 板 块 财 务 模 板', 'L', C_HOME,
          '💡 平时只录两张蓝色的表：【收支登记】（所有板块的收支都记这一张）和【应收应付登记】（赊销赊购）；'
          '各板块表、查询、资金台账、应收应付跟进、对账单都是自动的。所有表都可以插行、删行。')
    ws['A1'].value = '=IF(P_单位名称="","",P_单位名称&" · ")&"多板块收支与往来台账"'
    put(ws, 'B4', '截止日期', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, HOME_END, None, F_SEL, FILL_SEL, DATE, AC)
    dv_date(ws, HOME_END)
    put(ws, 'E4', '="现在算到："&TEXT(P_截止,"yyyy-mm-dd")&"（空着＝最后一笔 "&TEXT(P_最后日期,"yyyy-mm-dd")&"）；建账日 "&TEXT(P_建账日,"yyyy-mm-dd")',
        F_NOTE, align=AL, border=False)
    ws.merge_cells('E4:L4')
    put(ws, 'B5', '截止日期决定：各表默认算到哪天、应收应付的逾期天数从哪天算。各查看表也可以自己选起止日期。', F_NOTE, align=AL, border=False)
    ws.merge_cells('B5:L5')

    d, y0 = 'P_截止', 'P_年初'
    section(ws, 6, 'B', 'L', '关键数（截至截止日期；点名字跳过去）', C_HOME)
    kpis = [('资金余额（全部账户）', f'=SUM(账户_期初余额)+SUMIFS(收_净额,收_日期,"<="&{d})', SH_FUND),
            ('本年收入', f'=SUMIFS(收_板块收入,{dr("收_日期", y0, d)})', SH_QRY),
            ('本年支出', f'=SUMIFS(收_板块支出,{dr("收_日期", y0, d)})', SH_QRY),
            ('粮食存量（KG）', f'=SUM(板块_期初存量)+SUMIFS(收_入库量,收_日期,"<="&{d})-SUMIFS(收_出库量,收_日期,"<="&{d})'
                          f'+SUMIFS(往_入库量,往_日期,"<="&{d})-SUMIFS(往_出库量,往_日期,"<="&{d})', '粮食购销'),
            ('应收余额（别人欠我们）', f'=SUMIFS(往_应收额,往_日期,"<="&{d})-SUMIFS(收_冲应收,收_日期,"<="&{d})', SH_FOLLOW),
            ('其中已逾期', '=SUMIFS(往_未结,往_类型,"应收",往_逾期天数,">0")', SH_FOLLOW),
            ('应付余额（我们欠别人）', f'=SUMIFS(往_应付额,往_日期,"<="&{d})-SUMIFS(收_冲应付,收_日期,"<="&{d})', SH_FOLLOW),
            ('其中已逾期', '=SUMIFS(往_未结,往_类型,"应付",往_逾期天数,">0")', SH_FOLLOW)]
    for i, (lab, f, sh) in enumerate(kpis):
        r = 7 + (i // 4) * 2
        c1 = ['B', 'E', 'H', 'K'][i % 4]
        c2 = CL(CI(c1) + 1)
        link(put(ws, f'{c1}{r}', lab, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), fill('FFD9E1F2'), align=AC), sh)
        ws.merge_cells(f'{c1}{r}:{c2}{r}')
        put(ws, f'{c1}{r + 1}', f, F_KPI_V, fill('FFFFFFFF'), '#,##0.##' if 'KG' in lab else MONEY, AC)
        ws.merge_cells(f'{c1}{r + 1}:{c2}{r + 1}')
        ws.row_dimensions[r + 1].height = 26

    # 各板块一览
    r0 = 12
    section(ws, r0, 'B', 'L', '各板块一览（本年＝截止日期那年 1 月 1 日到截止日期）', C_VIEW)
    header(ws, r0 + 1, [('B', '业务板块'), ('C', '结余'), ('E', '本年收入'), ('F', '本年支出'), ('H', '应收余额'), ('I', '应付余额'),
                        ('K', '粮食存量KG'), ('L', '固定资产')], C_VIEW, height=24)
    for c in ('D', 'G', 'J'):
        put(ws, f'{c}{r0 + 1}', None, F_HDR, fill(C_VIEW))
    for i in range(N_SEG):
        r = r0 + 2 + i
        s = f'INDEX(板块_名称,{i + 1})'
        e = esc(s)
        g = lambda f: f'=IF({s}="","",{f})'
        put(ws, f'B{r}', f'=IF({s}="","",{s})', F_TXTB, align=AL)
        put(ws, f'C{r}', g(f'INDEX(板块_期初结余,{i + 1})+SUMIFS(收_板块收入,收_板块,{e},收_日期,"<="&{d})-SUMIFS(收_板块支出,收_板块,{e},收_日期,"<="&{d})'),
            F_AUTOB, fmt=MONEY, align=AR)
        put(ws, f'E{r}', g(f'SUMIFS(收_板块收入,收_板块,{e},{dr("收_日期", y0, d)})'), F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'F{r}', g(f'SUMIFS(收_板块支出,收_板块,{e},{dr("收_日期", y0, d)})'), F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'H{r}', g(f'SUMIFS(往_应收额,往_板块,{e},往_日期,"<="&{d})-SUMIFS(收_冲应收,收_板块,{e},收_日期,"<="&{d})'), F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'I{r}', g(f'SUMIFS(往_应付额,往_板块,{e},往_日期,"<="&{d})-SUMIFS(收_冲应付,收_板块,{e},收_日期,"<="&{d})'), F_AUTO, fmt=MONEY, align=AR)
        put(ws, f'K{r}', g(f'INDEX(板块_期初存量,{i + 1})+SUMIFS(收_入库量,收_板块,{e},收_日期,"<="&{d})-SUMIFS(收_出库量,收_板块,{e},收_日期,"<="&{d})'
                           f'+SUMIFS(往_入库量,往_板块,{e},往_日期,"<="&{d})-SUMIFS(往_出库量,往_板块,{e},往_日期,"<="&{d})'),
            F_AUTO, fmt='#,##0.##;[Red]-#,##0.##;"-"', align=AR)
        put(ws, f'L{r}', g(f'INDEX(板块_期初固定资产,{i + 1})+SUMIFS(收_固定资产,收_板块,{e},收_日期,"<="&{d})'
                           f'+SUMIFS(往_固定资产,往_板块,{e},往_日期,"<="&{d})'), F_AUTO, fmt=MONEY, align=AR)
        for c in ('D', 'G', 'J'):
            put(ws, f'{c}{r}', None)
    rt = r0 + 2 + N_SEG
    put(ws, f'B{rt}', '合计', F_TXTB, FILL_TOT, align=AC)
    for c in ('C', 'E', 'F', 'H', 'I', 'K', 'L'):
        put(ws, f'{c}{rt}', f'=SUM({c}{r0 + 2}:{c}{rt - 1})', F_TXTB, FILL_TOT, '#,##0.##' if c == 'K' else MONEY, AR)
    for c in ('D', 'G', 'J'):
        put(ws, f'{c}{rt}', None, fill_=FILL_TOT)
    # 没有名字的板块行不画框
    ws.conditional_formatting.add(f'B{r0 + 2}:L{rt - 1}', FormulaRule(formula=[f'$B{r0 + 2}=""'], border=Border()))

    # 提醒
    r = rt + 2
    section(ws, r, 'B', 'L', '提醒', C_CHK)
    CR = s_check.CHECK_ROWS
    rem = [f'="数据校验：要改 "&{SH_CHK}!{s_check.CHK_X}&" 项、要看 "&{SH_CHK}!{s_check.CHK_W}&" 项"&'
           f'IF({SH_CHK}!{s_check.CHK_X}+{SH_CHK}!{s_check.CHK_W}=0,"，都没问题 √","  → 点这里去看")',
           f'=IF(N({SH_CHK}!$C${CR[SH_CASH]})+N({SH_CHK}!$C${CR[SH_WL]})>0,"登记表有 "&({SH_CHK}!$C${CR[SH_CASH]}+{SH_CHK}!$C${CR[SH_WL]})&" 行要改（看那一行最右边的「这一行的问题」）","两张登记表都没有错 √")',
           '=IF(COUNTIFS(往_未结,">0",往_逾期天数,">0",往_类型,"应收")>0,COUNTIFS(往_未结,">0",往_逾期天数,">0",往_类型,"应收")&" 笔应收已经过了约定日期还没收回，共 "&TEXT(SUMIFS(往_未结,往_逾期天数,">0",往_类型,"应收"),"#,##0.00")&" → 看【应收应付跟进】","没有逾期的应收 √")',
           '=IF(COUNTIFS(往_未结,">0",往_逾期天数,">0",往_类型,"应付")>0,COUNTIFS(往_未结,">0",往_逾期天数,">0",往_类型,"应付")&" 笔应付已经过了约定日期还没付，共 "&TEXT(SUMIFS(往_未结,往_逾期天数,">0",往_类型,"应付"),"#,##0.00"),"")',
           f'=IF(ABS(N({SH_CHK}!$C${CR["内部转账没配对（一出一进合计不为 0）"]}))>=0.01,"内部转账有一笔只记了一边（差 "&TEXT({SH_CHK}!$C${CR["内部转账没配对（一出一进合计不为 0）"]},"#,##0.00")&"）","")']
    for i, f in enumerate(rem):
        rr = r + 1 + i
        put(ws, f'B{rr}', f, F_TXTB if i == 0 else F_TXT, align=AL, border=False)
        ws.merge_cells(f'B{rr}:L{rr}')
    link(ws[f'B{r + 1}'], SH_CHK)
    ws.conditional_formatting.add(f'B{r + 1}:B{r + len(rem)}', FormulaRule(
        formula=[f'OR(ISNUMBER(SEARCH("要改",B{r + 1})),ISNUMBER(SEARCH("逾期",B{r + 1})),ISNUMBER(SEARCH("只记了一边",B{r + 1})))'],
        font=Font(name=YH, sz=10, bold=True, color='FFC00000')))

    # 导航
    r = r + len(rem) + 2
    section(ws, r, 'B', 'L', '各张表（点名字跳过去；每张表右上角有「← 回首页」）', C_HOME)
    r += 1
    col_sets = [('B', 'C', 'F'), ('H', 'I', 'L')]
    for grp, color, sheets in NAV:
        put(ws, f'B{r}', grp, F_SEC, fill(color), align=AL)
        ws.merge_cells(f'B{r}:L{r}')
        r += 1
        for j, (sh, desc) in enumerate(sheets):
            c1, c2, c3 = col_sets[j % 2]
            rr = r + j // 2
            link(put(ws, f'{c1}{rr}', sh, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), align=AL), sh)
            put(ws, f'{c2}{rr}', desc, F_NOTE, align=AL)
            ws.merge_cells(f'{c2}{rr}:{c3}{rr}')
        r += (len(sheets) + 1) // 2

    # 怎么记
    r += 1
    section(ws, r, 'B', 'L', '常见情况怎么记', C_HOME)
    for i, (k, v) in enumerate(HOWTO):
        rr = r + 1 + i
        put(ws, f'B{rr}', k, F_TXTB, fill('FFF2F2F2'), align=ALW)
        ws.merge_cells(f'B{rr}:C{rr}')
        put(ws, f'E{rr}', v, F_TXT, align=ALW)
        ws.merge_cells(f'E{rr}:L{rr}')
        ws.row_dimensions[rr].height = 30
    r = r + len(HOWTO) + 2
    section(ws, r, 'B', 'L', '注意', C_HOME)
    notes = ['① 业务板块、收支项目、账户、经手人、往来单位都从下拉选；新的先到【基础资料】【往来单位】加一行。',
             '② 两张登记表哪里都可以插行、删行、排序（整行一起动）；绿色、橙色的表是自动的，有保护（没有密码），只有黄格子能选。',
             '③ 收回欠款、支付欠款一定要选往来单位，应收应付和对账单才会冲掉；业务板块选当初赊账的那个板块。',
             '④ 示例行（备注里写着「示例」）是照你原表的样子放的，正式用时整行删掉就行。']
    for i, t in enumerate(notes):
        rr = r + 1 + i
        put(ws, f'B{rr}', t, F_TXT, align=ALW, border=False)
        ws.merge_cells(f'B{rr}:L{rr}')
        ws.row_dimensions[rr].height = 20
    print_setup(ws, None, landscape=False)
