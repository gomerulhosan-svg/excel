# -*- coding: utf-8 -*-
"""查看表（绿）：资金日报、资金月报、费用统计。只用 layout 的定义名称和本表格子，不引用录入表。

【资金日报】黄格 C3 日期（空＝P_截止；晚于截止日照样按这天算，灰字提醒）。
   第 4 行（固定位置，首页可引用）：C4 当天可用资金、E4 手上票据、G4 当天实际收入、I4 当天实际支出（后两个不含内部转账）。
   ① 各账户（第 7 行起最多 15 个＋3 行合计）：昨日余额（到 d-1）、今日收入、今日支出（含内部转账）、今日余额、本月收入、本月支出（1 日～d）；
      合计：可用资金（账户_可用=1）、手上票据（账户_可用=0）、全公司当天实际收支（资_归类<>内部转账）。
   ② 当天每一笔（第 29 行起最多 80 笔）：按 资_排序键（同一天按录入顺序），本账户余额＝期初＋SUMIFS(净额, 账户, 排序键<=本笔)。
【资金月报】黄格 B3 年月（2026-10 / 202610；空＝P_截止年月）。
   第 4 行：B4 月末可用资金、D4 月末手上票据、F4 本月实际收入、H4 本月实际支出、J4 本月净现金流（后三个不含内部转账）。
   ① 各账户（第 7 行起）：月初余额（上月底，≤截止）、本月收入、支出（含内部转账，资_年月=m，≤截止）、月末余额（月底与截止日较早的）；合计：可用资金、手上票据。
   ② 全年各月（第 27～38 行＝1～12 月，第 39 行全年）：收客户货款、付供应商货款、费用支出、其他收入、买设备、不算收支（都按 资_年月）、
      本月净现金流（不含内部转账）、月末可用资金、月末手上票据（截止日以后的月份空着）。
   ③ 按收支类别（第 44 行起，按 类别_名称 顺序）：算到哪、本月收入、支出、本年累计收入、支出；没归类的一行、合计（不含内部转账）、其中内部转账。
   （变长的清单放最后，中间不留大段空行；所以「全年各月」排第二、「按收支类别」排第三。）
【费用统计】黄格 B3 年份（空＝P_年度）。第 4 行：B4 全年费用合计，E4/G4/I4/K4/M4/O4 六类全年小计。
   ① 左边 A～O：6 类费用（EXPENSE_KINDS 顺序）× 1～12 月＋全年＋占比，每类列出 类别_归类=这类 的类别（最多 20 个），段尾小计，最后合计。
      金额＝SUMIFS(资_支出额, 资_类别, 类别, 资_所属年月, m, 资_日期, "<="&P_截止)（按所属月份）。
   ② 右边 Q～W：黄格 R3 类别（空＝全部费用）、U3 月份（1～12，空＝全年）：这一年（按所属月份）每一笔，最多 150 笔，按日期。
      隐藏列 BX/BY（第 2 行起，每条资金一行）挑出符合条件的排序键，再 SMALL。
隐藏列：各表 AA 起（AA 名字、AB 值是标量）；counter() 的条件列从 CA 往右。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from layout import *
from common import *

M = SORT_M
GREEN_H = 'FF70AD47'
_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
F_KV = Font(name=YH, sz=10, bold=True, color='FFC00000')
LBL = fill('FFD9E1F2')
MONEY_B = '#,##0.00;[Red]-#,##0.00;""'
YM_FMT = '[<100000]yyyy-mm;0'
DAY_SHOW, MON_CAT_SHOW, FEE_SEG_MAX, FEE_SHOW = 80, N_CAT, 20, 150


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def R2(x):
    return f'ROUND({x},2)'


class Sc:
    """隐藏标量：AA 列名字、AB 列值（第 r0 行起）"""

    def __init__(self, ws, keys, r0=3, cl='AA', cv='AB'):
        self.ws, self.cv = ws, cv
        self.r = {k: r0 + i for i, k in enumerate(keys)}
        for k, r in self.r.items():
            ws[f'{cl}{r}'] = k
            ws[f'{cl}{r}'].font = F_HELP

    def __call__(self, k):
        return f'${self.cv}${self.r[k]}'

    def set(self, k, f):
        c = self.ws[f'{self.cv}{self.r[k]}']
        c.value = '=' + f
        c.font = F_HELP


def hset(ws, coord, f):
    ws[coord] = f if isinstance(f, (int, float)) else '=' + f
    ws[coord].font = F_HELP


def cell(ws, coord, f, font=F_TXT, fmt=None, align=AC):
    c = ws[coord]
    c.value = '=' + f
    c.font = font
    if fmt:
        c.number_format = fmt
    c.alignment = align
    return c


def note(ws, coord, last, f, height=None):
    r = coord[1:] if coord[1].isdigit() else coord[2:]
    ws.merge_cells(f'{coord}:{last}{r}')
    put(ws, coord, f, F_NOTE, align=ALW, border=False)
    if height:
        ws.row_dimensions[int(r)].height = height


def kpi(ws, lbl_cell, lbl, val_cell, f, fmt=MONEY, merge_to=None):
    put(ws, lbl_cell, lbl, F_KPI_L, LBL, align=AC)
    put(ws, val_cell, '=' + f, F_KV, FILL_AUTO, fmt, AR)
    if merge_to:
        ws.merge_cells(f'{val_cell}:{merge_to}')


def pick(idx, col, r0, n, choices):
    """idx>0：取隐藏表第 idx 行；idx<0：CHOOSE(-idx, 合计行的值…)；0：空"""
    return (f'IF({idx}=0,"",IF({idx}>0,INDEX(${col}${r0}:${col}${r0 + n - 1},{idx}),'
            f'CHOOSE(-{idx},{",".join(choices)})))')


def month_any(x):
    """真日期 → 年月；202610 → 本身；看不懂 → 0"""
    return (f'IF(AND({x}>=200001,{x}<=209912,INT({x})={x},MOD({x},100)>=1,MOD({x},100)<=12),{x},'
            f'IF(AND({x}>=36526,{x}<73051),YEAR({x})*100+MONTH({x}),0))')


def mhead(ws, row, c1, c2):
    ws.merge_cells(f'{c1}{row}:{c2}{row}')
    for i in range(CI(c1), CI(c2) + 1):
        ws.cell(row=row, column=i).border = BD


def finish(ws, last, hdr_rows, end_cell, cap_row, freeze='A5'):
    ws.freeze_panes = freeze
    print_setup(ws, hdr_rows, landscape=True)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${last}$1:${last}${cap_row},{q}!{end_cell})')


def acct_table(ws, r0, vals):
    """隐藏：账户第 a 个（a＝1..N_ACC，第 r0+a-1 行）：AD 名称、AE 类型、AF 可用、AG 票据、AN 显示（非空且第一次出现）；
       vals：{列: 公式(a, r)}；AO 计数列（counter）"""
    for i in range(N_ACC):
        a, r = i + 1, r0 + i
        hset(ws, f'AD{r}', f'INDEX(账户_名称,{a})&""')
        hset(ws, f'AN{r}', f'IF(AD{r}="",0,IF(MATCH({esc(f"AD{r}")},账户_名称,0)={a},1,0))')
        hset(ws, f'AE{r}', f'IF(AN{r}=1,INDEX(账户_类型,{a}),"")')
        hset(ws, f'AF{r}', f'IF(AN{r}=1,INDEX(账户_可用,{a}),0)')
        hset(ws, f'AG{r}', f'IF(AND(AN{r}=1,INDEX(账户_可用,{a})=0),1,0)')
        for col, fn in vals.items():
            hset(ws, f'{col}{r}', f'IF(AN{r}<>1,0,{fn(a, r)})')
    counter(ws, 'AO', r0, N_ACC, lambda i: f'$AN{r0 + i}=1')


def tot(col, flag, r0):
    return R2(f'SUMIFS(${col}${r0}:${col}${r0 + N_ACC - 1},${flag}${r0}:${flag}${r0 + N_ACC - 1},1)')


# ═══════════════════════════════ 资金日报 ═══════════════════════════════
TIP_DAY = ('💡 每天看一眼钱。黄格选日期（空着＝截止日，截止日在首页改）。① 每个账户：昨日余额、今天收了多少、付了多少、今日余额，'
           '还有这个月 1 号到这天的收支（都含内部转账，余额才接得上）；下面合计：可用资金（银行、微信、现金……）、手上票据（承兑汇票，不能直接花）、'
           '全公司当天实际收支（不含自己账户之间倒钱）。② 当天每一笔，同一天按录入顺序，带这一笔之后本账户的余额。'
           '余额一律按日期算（资金台帐右边的即时余额按表里上下顺序算，日期乱了两边会差一点，按日期排一下序就一样了）。')

D_A0, D_N1 = 7, N_ACC + 3          # ① 第 7 行起 18 行
D_NOTE1 = D_A0 + D_N1              # 25
D_SEC2, D_NOTE2, D_H2 = D_NOTE1 + 1, D_NOTE1 + 2, D_NOTE1 + 3
D_L0 = D_H2 + 1                    # 29
D_NL = DAY_SHOW + 2                # 80 笔＋合计＋共几笔


def build_day(ws):
    LAST = 'I'
    widths(ws, {'A': 5.5, 'B': 16, 'C': 15, 'D': 14, 'E': 24, 'F': 13.5, 'G': 13.5, 'H': 14, 'I': 13.5})
    title(ws, '=IF(P_公司名称="","",P_公司名称&"　")&"资金日报"', LAST, C_VIEW, TIP_DAY)
    selector(ws, 'B3', '选日期', 'C3', None, fmt=DATE)
    dv_date(ws, 'C3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 30

    keys = ['d', '认出', '月初', '账户数', '笔数', '前面', '显示数', '末行', '内转净',
            '全_今收', '全_今支', '全_月收', '全_月支', '全天收入', '全天支出']
    for p in ('可用', '票据'):
        keys += [f'{p}_{k}' for k in ('昨', '今收', '今支', '今余', '月收', '月支')]
    sc = Sc(ws, keys)
    d, m0 = sc('d'), sc('月初')
    sc.set('d', f'IF(AND(ISNUMBER(C3),C3>=36526,C3<73051),INT(C3),P_截止)')
    sc.set('认出', 'IF(TRIM(C3&"")="",1,IF(AND(ISNUMBER(C3),C3>=36526,C3<73051),1,0))')
    sc.set('月初', f'DATE(YEAR({d}),MONTH({d}),1)')
    sc.set('笔数', f'COUNTIFS(资_排序键,">="&{d}*{M},资_排序键,"<"&({d}+1)*{M})')
    sc.set('前面', f'COUNTIF(资_排序键,"<"&{d}*{M})')
    sc.set('显示数', f'MIN({sc("笔数")},{DAY_SHOW})')
    sc.set('末行', f'{D_H2}+{sc("显示数")}+2')
    sc.set('内转净', R2(f'SUMIFS(资_净额,资_归类,"内部转账",资_日期,{d})'))
    real = '资_资金有效,1,资_归类,"<>内部转账"'
    mon = f'资_日期,">="&{m0},资_日期,"<="&{d}'
    sc.set('全_今收', R2(f'SUMIFS(资_收入,资_日期,{d},{real})'))
    sc.set('全_今支', R2(f'SUMIFS(资_支出,资_日期,{d},{real})'))
    sc.set('全_月收', R2(f'SUMIFS(资_收入,{mon},{real})'))
    sc.set('全_月支', R2(f'SUMIFS(资_支出,{mon},{real})'))
    sc.set('全天收入', R2(f'SUMIFS(资_收入,资_日期,{d},资_资金有效,1)'))
    sc.set('全天支出', R2(f'SUMIFS(资_支出,资_日期,{d},资_资金有效,1)'))

    # ── 隐藏账户表（第 7 行起，a＝1..15） ──
    H = dict(昨='AH', 今收='AI', 今支='AJ', 今余='AK', 月收='AL', 月支='AM')
    acct_table(ws, D_A0, {
        'AH': lambda a, r: R2(f'INDEX(账户_期初余额,{a})+SUMIFS(资_净额,资_账户号,{a},资_日期,"<"&{d})'),
        'AI': lambda a, r: R2(f'SUMIFS(资_收入,资_账户号,{a},资_日期,{d},资_资金有效,1)'),
        'AJ': lambda a, r: R2(f'SUMIFS(资_支出,资_账户号,{a},资_日期,{d},资_资金有效,1)'),
        'AK': lambda a, r: R2(f'AH{r}+AI{r}-AJ{r}'),
        'AL': lambda a, r: R2(f'SUMIFS(资_收入,资_账户号,{a},资_资金有效,1,{mon})'),
        'AM': lambda a, r: R2(f'SUMIFS(资_支出,资_账户号,{a},资_资金有效,1,{mon})'),
    })
    sc.set('账户数', f'MIN({cnt("AO", D_A0, N_ACC)},{N_ACC})')
    for p, flag in (('可用', 'AF'), ('票据', 'AG')):
        for k, col in H.items():
            sc.set(f'{p}_{k}', tot(col, flag, D_A0))

    # ── 第 3、4 行：实际用的、要紧的数 ──
    ws.merge_cells('D3:H3')
    put(ws, 'D3', (f'="实际用的："&TEXT({d},"yyyy-mm-dd")&IF(TRIM(C3&"")="","（空着＝截止日）","")'
                   f'&IF({sc("认出")}=0,"（⚠ 没认出这个日期，先按截止日）","")'
                   f'&IF({d}>P_截止,"　⚠ 晚于截止日 "&TEXT(P_截止,"yyyy-mm-dd")&"：截止日以后的不算（月报、利润表都不算），日报照样按这一天算","")'
                   f'&IF({d}<P_建账日,"　早于建账日 "&TEXT(P_建账日,"yyyy-mm-dd")&"：建账前没有流水，余额都是期初余额","")'),
        F_NOTE, align=ALW, border=False)
    kpi(ws, 'B4', '当天可用资金', 'C4', sc('可用_今余'))
    kpi(ws, 'D4', '手上票据', 'E4', sc('票据_今余'))
    kpi(ws, 'F4', '当天实际收入', 'G4', sc('全_今收'))
    kpi(ws, 'H4', '当天实际支出', 'I4', sc('全_今支'))
    ws.row_dimensions[4].height = 22

    # ── ① 各账户 ──
    section(ws, 5, 'A', LAST, '① 各账户（余额按日期算；收入、支出含内部转账，余额才接得上）', C_VIEW)
    header(ws, 6, [('A', '序号'), ('B', '账户'), ('C', '类型'), ('D', '昨日余额'), ('E', '今日收入'), ('F', '今日支出'),
                   ('G', '今日余额'), ('H', '本月收入\n（1 日～当天）'), ('I', '本月支出\n（1 日～当天）')], GREEN_H, height=34)
    na = sc('账户数')
    for k in range(1, D_N1 + 1):
        r = D_A0 + k - 1
        ix = f'$AP{r}'
        hset(ws, f'AP{r}', f'IF({k}<={na},{kth(k, "AO", D_A0, N_ACC)},IF({k}={na}+1,-1,IF({k}={na}+2,-2,IF({k}={na}+3,-3,0))))')
        cell(ws, f'A{r}', f'IF({ix}>0,{k},"")', F_NOTE)
        cell(ws, f'B{r}', pick(ix, 'AD', D_A0, N_ACC, ['"可用资金合计"', '"手上票据"', '"全公司实际收支"']), align=AL)
        cell(ws, f'C{r}', pick(ix, 'AE', D_A0, N_ACC, ['"不含票据"', '"承兑汇票"', '"不含内部转账"']), align=AL)
        for col, k_ in zip('DEFGHI', H):
            ch = [sc(f'可用_{k_}'), sc(f'票据_{k_}'), sc(f'全_{k_}') if k_ in ('今收', '今支', '月收', '月支') else '""']
            cell(ws, f'{col}{r}', pick(ix, H[k_], D_A0, N_ACC, ch), fmt=MONEY, align=AR)
    a0, a1 = D_A0, D_A0 + D_N1 - 1
    cf = ws.conditional_formatting
    z0 = f'$AP{a0}'
    rng = f'A{a0}:{LAST}{a1}'
    cf.add(rng, FormulaRule(formula=[f'OR({z0}=-1,{z0}=-2)'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD,
                            stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}=-3'], fill=cf_fill('FFDDEBF7'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}>0'], border=CF_BD))
    note(ws, f'A{D_NOTE1}', LAST,
         f'=IF({sc("内转净")}<>0,"⚠ 当天的内部转账转出、转进差 "&TEXT(ABS({sc("内转净")}),"#,##0.00")&"（只记了一边？手续费要另记一行）　","")'
         f'&"可用资金＝银行、微信、支付宝、现金等账户的合计；手上票据＝承兑汇票账户（基础资料里类型是「票据」的）。'
         f'全公司实际收支不算内部转账（自己账户之间倒钱、承兑贴现转进银行）。"', height=30)
    cf.add(f'A{D_NOTE1}', FormulaRule(formula=[f'LEFT($A${D_NOTE1},1)="⚠"'], font=Font(bold=True, color='FFC00000')))

    # ── ② 当天每一笔 ──
    section(ws, D_SEC2, 'A', LAST, '② 当天每一笔（同一天按录入顺序；余额＝这一笔之后本账户的余额）', C_VIEW)
    N, NN = sc('笔数'), sc('显示数')
    note(ws, f'A{D_NOTE2}', LAST,
         f'=IF({N}=0,TEXT({d},"yyyy-mm-dd")&" 没有收支",TEXT({d},"yyyy-mm-dd")&" 共 "&{N}&" 笔（含内部转账）")'
         f'&IF({N}>{DAY_SHOW},"　⚠ 只列前 {DAY_SHOW} 笔（合计是全天的）","")')
    cf.add(f'A{D_NOTE2}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",$A${D_NOTE2}))'], font=Font(bold=True, color='FFC00000')))
    header(ws, D_H2, [('A', '序号'), ('B', '账户'), ('C', '收支类别'), ('D', '往来单位'), ('E', '摘要'), ('F', '收入'),
                      ('G', '支出'), ('H', '本账户余额'), ('I', '算到哪')], GREEN_H, height=30)
    for k in range(1, D_NL + 1):
        r = D_L0 + k - 1
        q, y, t, a = f'$AQ{r}', f'$AR{r}', f'$AS{r}', f'$AU{r}'
        ws[f'AQ{r}'] = k
        ws[f'AQ{r}'].font = F_HELP
        hset(ws, f'AR{r}', f'IF({q}>{NN},"",SMALL(资_排序键,{sc("前面")}+{q}))')
        hset(ws, f'AT{r}', f'IF({y}="",0,MOD({y},{M}))')
        n = f'$AT{r}'
        hset(ws, f'AS{r}', f'IF({n}>0,1,IF({q}={NN}+1,3,IF({q}={NN}+2,4,0)))')
        hset(ws, f'AU{r}', f'IF({t}=1,INDEX(资_账户号,{n}),0)')
        ix = lambda nm: f'INDEX({nm},{n})'
        cell(ws, f'A{r}', f'IF({t}=1,{q},"")', F_NOTE)
        cell(ws, f'B{r}', f'IF({t}=1,{ix("资_账户")},IF({t}=3,"当天合计",IF({t}=4,"共 "&{N}&" 笔","")))', align=AL)
        cell(ws, f'C{r}', f'IF({t}=1,{ix("资_类别")},IF({t}=3,"含内部转账",""))', align=AL)
        cell(ws, f'D{r}', f'IF({t}=1,{ix("资_单位")},"")', align=AL)
        cell(ws, f'E{r}', f'IF({t}=1,{ix("资_显示摘要")},IF(AND({t}=4,{N}>{DAY_SHOW}),"⚠ 只列了前 {DAY_SHOW} 笔（合计是全天的）",""))',
             align=AL)
        cell(ws, f'F{r}', f'IF({t}=1,IF({ix("资_收入")}=0,"",{ix("资_收入")}),IF({t}=3,{sc("全天收入")},""))', fmt=MONEY_B, align=AR)
        cell(ws, f'G{r}', f'IF({t}=1,IF({ix("资_支出")}=0,"",{ix("资_支出")}),IF({t}=3,{sc("全天支出")},""))', fmt=MONEY_B, align=AR)
        cell(ws, f'H{r}', f'IF({t}=1,ROUND(INDEX(账户_期初余额,{a})+SUMIFS(资_净额,资_账户号,{a},资_排序键,"<="&{y}),2),"")',
             fmt=MONEY, align=AR)
        cell(ws, f'I{r}', f'IF({t}=1,IF({ix("资_归类")}="","？没归类",{ix("资_归类")}),"")', align=AC)
    d0, d1 = D_L0, D_L0 + D_NL - 1
    t0 = f'$AS{d0}'
    rng = f'A{d0}:{LAST}{d1}'
    cf.add(rng, FormulaRule(formula=[f'{t0}=3'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=4'], font=Font(color='FF808080'), stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'AND({t0}=1,$I{d0}="内部转账")'], fill=cf_fill('FFEAF4FB'), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=1'], border=CF_BD))
    cf.add(f'E{d0}:E{d1}', FormulaRule(formula=[f'AND({t0}=4,LEFT($E{d0},1)="⚠")'], font=Font(bold=True, color='FFC00000')))
    hide(ws, *[CL(i) for i in range(CI('J'), CI('AU') + 1)])
    finish(ws, LAST, f'{D_H2}:{D_H2}', sc('末行'), d1)


# ═══════════════════════════════ 资金月报 ═══════════════════════════════
TIP_MON = ('💡 黄格选年月（填 2026-10 或 202610 都行；空着＝截止日那个月）。只算截止日以前的钱。'
           '① 每个账户：月初余额（上个月底）、本月收入、本月支出（含内部转账）、月末余额（月底，没到月底按截止日）。'
           '② 这一年 1～12 月：收客户货款、付供应商货款、费用支出、其他收入、买设备、不算收支（借款还款、老板存取、保证金等）、'
           '本月净现金流（不含内部转账）、月末可用资金和手上票据。'
           '③ 按收支类别：本月、今年 1 月到本月累计的收入和支出；合计不含内部转账，内部转账单列一行，看转出、转进对不对得上。')

MO_A0, MO_N1 = 7, N_ACC + 2         # ① 第 7 行起 17 行
MO_NOTE1 = MO_A0 + MO_N1            # 24
MO_SEC2, MO_H2 = MO_NOTE1 + 1, MO_NOTE1 + 2   # 25、26
MO_M0 = MO_H2 + 1                   # 27～38 月，39 全年
MO_NOTE2 = MO_M0 + 13               # 40
MO_SEC3, MO_NOTE3, MO_H3 = MO_NOTE2 + 1, MO_NOTE2 + 2, MO_NOTE2 + 3   # 41、42、43
MO_C0 = MO_H3 + 1                   # 44
MO_NC = N_CAT + 3                   # 类别＋没归类＋合计＋内部转账
MAT0 = 'BA'                         # ② 隐藏：月末余额矩阵（账户 1..15 横排），第 MO_SEC2 行＝可用标志、MO_H2 行＝票据标志


def build_mon(ws):
    LAST = 'J'
    widths(ws, dict({'A': 12}, **{c: 13.5 for c in 'BCDEFGHIJ'}))
    title(ws, '=IF(P_公司名称="","",P_公司名称&"　")&"资金月报"', LAST, C_VIEW, TIP_MON)
    selector(ws, 'A3', '选年月', 'B3', None, fmt=YM_FMT)
    dv = DataValidation(type='custom', formula1='TRUE', allow_blank=True, showErrorMessage=False, showInputMessage=True,
                        promptTitle='提示', prompt='填 2026-10 或 202610；空着＝截止日那个月')
    ws.add_data_validation(dv)
    dv.add('B3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 30

    keys = ['输入', '认出前', 'm', '认出', 'Y', 'mo', '月初', '月末', '起日', '末日', '年初月', '账户数', '类别数', '有没归类',
            '显示行数', '末行', '全_收', '全_支', '全_净']
    for p in ('可用', '票据'):
        keys += [f'{p}_{k}' for k in ('初', '收', '支', '末')]
    for p in ('合计', '内转', '全部', '没归类'):
        keys += [f'{p}_{k}' for k in ('月收', '月支', '年收', '年支')]
    sc = Sc(ws, keys)
    v, m, Y, mo = sc('输入'), sc('m'), sc('Y'), sc('mo')
    sc.set('输入', 'IF(ISNUMBER(B3),B3,IFERROR(--TRIM(B3&""),0))')
    sc.set('认出前', month_any(v))
    sc.set('m', f'IF(TRIM(B3&"")="",P_截止年月,IF({sc("认出前")}=0,P_截止年月,{sc("认出前")}))')
    sc.set('认出', f'IF(TRIM(B3&"")="",1,IF({sc("认出前")}>0,1,0))')
    sc.set('Y', f'INT({m}/100)')
    sc.set('mo', f'MOD({m},100)')
    sc.set('月初', f'DATE({Y},{mo},1)')
    sc.set('月末', f'DATE({Y},{mo}+1,0)')
    sc.set('起日', f'MIN({sc("月初")}-1,P_截止)')
    sc.set('末日', f'MIN({sc("月末")},P_截止)')
    sc.set('年初月', f'{Y}*100+1')
    upto = '资_日期,"<="&P_截止'
    inm = f'资_资金有效,1,资_年月,{m},{upto}'
    iny = f'资_资金有效,1,资_年月,">="&{sc("年初月")},资_年月,"<="&{m},{upto}'
    sc.set('全_收', R2(f'SUMIFS(资_收入,{inm},资_归类,"<>内部转账")'))
    sc.set('全_支', R2(f'SUMIFS(资_支出,{inm},资_归类,"<>内部转账")'))
    sc.set('全_净', R2(f'SUMIFS(资_净额,资_年月,{m},{upto},资_归类,"<>内部转账")'))

    # ── 隐藏账户表 ──
    acct_table(ws, MO_A0, {
        'AH': lambda a, r: R2(f'INDEX(账户_期初余额,{a})+SUMIFS(资_净额,资_账户号,{a},资_日期,"<="&{sc("起日")})'),
        'AI': lambda a, r: R2(f'SUMIFS(资_收入,资_账户号,{a},{inm})'),
        'AJ': lambda a, r: R2(f'SUMIFS(资_支出,资_账户号,{a},{inm})'),
        'AK': lambda a, r: R2(f'INDEX(账户_期初余额,{a})+SUMIFS(资_净额,资_账户号,{a},资_日期,"<="&{sc("末日")})'),
    })
    sc.set('账户数', f'MIN({cnt("AO", MO_A0, N_ACC)},{N_ACC})')
    H = dict(初='AH', 收='AI', 支='AJ', 末='AK')
    for p, flag in (('可用', 'AF'), ('票据', 'AG')):
        for k, col in H.items():
            sc.set(f'{p}_{k}', tot(col, flag, MO_A0))

    ws.merge_cells('C3:I3')
    put(ws, 'C3', (f'="＝ "&{Y}&" 年 "&{mo}&" 月"&IF(TRIM(B3&"")="","（空着＝截止日那个月）","")'
                   f'&IF({sc("认出")}=0,"（⚠ 没认出来，先按截止日那个月）","")'
                   f'&"；月初余额＝"&TEXT({sc("月初")}-1,"yyyy-mm-dd")&"，月末余额＝"&TEXT({sc("末日")},"yyyy-mm-dd")'
                   f'&IF({sc("月末")}>P_截止,"（截止日；截止日以后的不算）","")'
                   f'&IF({sc("月初")}>P_截止,"　⚠ 这个月在截止日 "&TEXT(P_截止,"yyyy-mm-dd")&" 以后，还没有数","")'
                   f'&IF({sc("月末")}<P_建账日,"　⚠ 这个月在建账日以前，没有流水","")'),
        F_NOTE, align=ALW, border=False)
    kpi(ws, 'A4', '月末可用资金', 'B4', sc('可用_末'))
    kpi(ws, 'C4', '月末手上票据', 'D4', sc('票据_末'))
    kpi(ws, 'E4', '本月实际收入', 'F4', sc('全_收'))
    kpi(ws, 'G4', '本月实际支出', 'H4', sc('全_支'))
    kpi(ws, 'I4', '本月净现金流', 'J4', sc('全_净'))
    ws.row_dimensions[4].height = 22
    cf = ws.conditional_formatting

    # ── ① 各账户 ──
    section(ws, 5, 'A', LAST, '① 各账户（收入、支出含内部转账；月末＝月底和截止日较早的那天）', C_VIEW)
    header(ws, 6, [('A', '账户'), ('C', '类型'), ('D', '月初余额\n（上月底）'), ('E', '本月收入'), ('F', '本月支出'),
                   ('G', '月末余额'), ('H', '说明')], GREEN_H, height=34)
    mhead(ws, 6, 'A', 'B')
    mhead(ws, 6, 'H', LAST)
    na = sc('账户数')
    for k in range(1, MO_N1 + 1):
        r = MO_A0 + k - 1
        ix = f'$AP{r}'
        hset(ws, f'AP{r}', f'IF({k}<={na},{kth(k, "AO", MO_A0, N_ACC)},IF({k}={na}+1,-1,IF({k}={na}+2,-2,0)))')
        ws.merge_cells(f'A{r}:B{r}')
        cell(ws, f'A{r}', pick(ix, 'AD', MO_A0, N_ACC, ['"可用资金合计"', '"手上票据"']), align=AL)
        cell(ws, f'C{r}', pick(ix, 'AE', MO_A0, N_ACC, ['"不含票据"', '"承兑汇票"']), align=AL)
        for col, k_ in zip('DEFG', H):
            cell(ws, f'{col}{r}', pick(ix, H[k_], MO_A0, N_ACC, [sc(f'可用_{k_}'), sc(f'票据_{k_}')]), fmt=MONEY, align=AR)
        ws.merge_cells(f'H{r}:{LAST}{r}')
        cell(ws, f'H{r}', f'IF({ix}=0,"",IF(N(G{r})<-0.005,"⚠ 余额是负数：漏记了收入，还是记错了账户？",'
                          f'IF(ROUND(N(D{r})+N(E{r})-N(F{r})-N(G{r}),2)<>0,"⚠ 月初＋收入－支出≠月末","")))', F_NOTE, align=AL)
    a0, a1 = MO_A0, MO_A0 + MO_N1 - 1
    z0 = f'$AP{a0}'
    rng = f'A{a0}:{LAST}{a1}'
    cf.add(f'H{a0}:{LAST}{a1}', FormulaRule(formula=[f'LEFT($H{a0},1)="⚠"'], font=Font(bold=True, color='FFC00000')))
    cf.add(rng, FormulaRule(formula=[f'{z0}<0'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}>0'], border=CF_BD))
    note(ws, f'A{MO_NOTE1}', LAST, '="可用资金＝银行、微信、支付宝、现金等账户；手上票据＝承兑汇票账户（基础资料里类型是「票据」的）。'
                                   '各账户收支含内部转账（自己账户之间倒钱），全公司合起来一进一出抵掉。"')

    # ── ② 全年各月 ──
    section(ws, MO_SEC2, 'A', LAST, f'="② "&{Y}&" 年各月（按付款日期；月末＝月底和截止日较早的；截止日以后的月份空着）"', C_VIEW)
    header(ws, MO_H2, [('A', '月份'), ('B', '收客户货款\n（净收）'), ('C', '付供应商货款\n（净付）'), ('D', '费用支出\n（6 类费用）'),
                       ('E', '其他收入'), ('F', '买设备'), ('G', '不算收支\n（净额）'), ('H', '本月净现金流\n（不含内部转账）'),
                       ('I', '月末可用资金'), ('J', '月末手上票据')], GREEN_H, height=34)
    mat = [CL(CI(MAT0) + i) for i in range(N_ACC)]
    for i, c in enumerate(mat):
        r_ = MO_A0 + i
        hset(ws, f'{c}{MO_SEC2}', f'$AF${r_}')        # 可用标志（按账户顺序横排）
        hset(ws, f'{c}{MO_H2}', f'$AG${r_}')          # 票据标志
    fu_rng_a, fu_rng_b = f'${mat[0]}${MO_SEC2}:${mat[-1]}${MO_SEC2}', f'${mat[0]}${MO_H2}:${mat[-1]}${MO_H2}'
    for k in range(1, 14):
        r = MO_M0 + k - 1
        yr = k == 13
        # AR 本月键、AS 本月第一天、AT 月末（和截止日较早的）、AU 有余额、AV 截止日以后
        if yr:
            hset(ws, f'AS{r}', f'DATE({Y},1,1)')
            hset(ws, f'AT{r}', f'MIN(DATE({Y},12,31),P_截止)')
            hset(ws, f'AU{r}', f'IF(OR(DATE({Y},1,1)>P_截止,DATE({Y},12,31)<P_建账日-1),0,1)')
        else:
            hset(ws, f'AR{r}', f'{Y}*100+{k}')
            hset(ws, f'AS{r}', f'DATE({Y},{k},1)')
            hset(ws, f'AT{r}', f'MIN(DATE({Y},{k}+1,0),P_截止)')
            hset(ws, f'AU{r}', f'IF(OR(DATE({Y},{k},1)>P_截止,DATE({Y},{k}+1,0)<P_建账日-1),0,1)')
            hset(ws, f'AV{r}', f'IF(DATE({Y},{k},1)>P_截止,1,0)')
        for i, c in enumerate(mat):
            hset(ws, f'{c}{r}', f'IF($AU{r}=0,0,ROUND(INDEX(账户_期初余额,{i + 1})+SUMIFS(资_净额,资_账户号,{i + 1},资_日期,"<="&$AT{r}),2))')
        put(ws, f'A{r}', '全年' if yr else f'{k} 月', F_TXTB if yr else F_TXT, align=AC)
        if yr:
            for col in 'BCDEFGH':
                cell(ws, f'{col}{r}', f'ROUND(SUM({col}{MO_M0}:{col}{MO_M0 + 11}),2)', F_TXTB, MONEY, AR)
        else:
            mm, fu = f'$AR{r}', f'$AV{r}'
            w = f'资_年月,{mm},{upto}'
            fx = {
                'B': f'SUMIFS(资_净额,资_归类,"收客户货款",{w})',
                'C': f'SUMIFS(资_支出额,资_归类,"付供应商货款",{w})',
                'D': '+'.join(f'SUMIFS(资_支出额,资_归类,"{kd}",{w})' for kd in EXPENSE_KINDS),
                'E': f'SUMIFS(资_净额,资_归类,"其他收入",{w})',
                'F': f'SUMIFS(资_支出额,资_归类,"买设备",{w})',
                'G': f'SUMIFS(资_净额,资_归类,"不算收支",{w})',
                'H': f'SUMIFS(资_净额,{w},资_归类,"<>内部转账")',
            }
            for col, f in fx.items():
                cell(ws, f'{col}{r}', f'IF({fu}=1,"",{R2(f)})', F_TXTB if col == 'H' else F_TXT, MONEY, AR)
        mr = f'${mat[0]}{r}:${mat[-1]}{r}'
        bold = F_TXTB if yr else F_TXT
        cell(ws, f'I{r}', f'IF($AU{r}=0,"",ROUND(SUMPRODUCT({mr},{fu_rng_a}),2))', bold, MONEY, AR)
        cell(ws, f'J{r}', f'IF($AU{r}=0,"",ROUND(SUMPRODUCT({mr},{fu_rng_b}),2))', bold, MONEY, AR)
        for col in 'ABCDEFGHIJ':
            ws[f'{col}{r}'].border = BD
            if yr:
                ws[f'{col}{r}'].fill = FILL_TOT
    m0_, m1_ = MO_M0, MO_M0 + 11
    cf.add(f'A{m0_}:{LAST}{m1_}', FormulaRule(formula=[f'$AR{m0_}={m}'], fill=cf_fill('FFFFF2CC'), font=Font(bold=True)))
    cf.add(f'A{m0_}:{LAST}{m1_}', FormulaRule(formula=[f'$AV{m0_}=1'], font=Font(color='FFA6A6A6')))
    note(ws, f'A{MO_NOTE2}', LAST,
         '="收客户货款、付供应商货款都是净额（退款已经减掉）；费用支出＝生产成本＋销售费用＋管理费用＋财务费用＋税金＋其他支出（按付款月份，'
         '利润表、费用统计按所属月份）；不算收支＝借款、还借款、老板存取、保证金、还信用卡的净额（正数＝净进来）。'
         '本月净现金流＝全部收入－全部支出（不含内部转账）；全年那行的月末＝年底或截止日。选的月份黄底。"', height=44)

    # ── ③ 按收支类别 ──
    section(ws, MO_SEC3, 'A', LAST, f'="③ 按收支类别（"&{Y}&" 年 "&{mo}&" 月；本年累计＝1 月～"&{mo}&" 月；只算截止日以前的）"', C_VIEW)
    note(ws, f'A{MO_NOTE3}', LAST,
         f'="按【基础资料】③收支类别的顺序。收入、支出是原数（退款记在哪栏就算哪栏）。合计不含内部转账；内部转账单列一行，转出、转进应该相等。"'
         f'&IF({sc("有没归类")}=1,"　⚠ 有些钱的收支类别没选或不在清单里，单列一行（到资金台帐看最右边的校验列）","")')
    cf.add(f'A{MO_NOTE3}', FormulaRule(formula=[f'ISNUMBER(FIND("⚠",$A${MO_NOTE3}))'], font=Font(bold=True, color='FFC00000')))
    header(ws, MO_H3, [('A', '收支类别'), ('C', '算到哪'), ('D', '本月收入'), ('E', '本月支出'), ('F', '本年累计收入\n（1 月～本月）'),
                       ('G', '本年累计支出\n（1 月～本月）'), ('H', '说明')], GREEN_H, height=34)
    mhead(ws, MO_H3, 'A', 'B')
    mhead(ws, MO_H3, 'H', LAST)
    # 隐藏类别表（第 MO_C0 行起，i＝1..N_CAT）：AD 名称、AN 显示、AE 归类、AH..AK 本月收/支、本年收/支
    C0 = MO_C0
    for i in range(N_CAT):
        p, r = i + 1, C0 + i
        hset(ws, f'AD{r}', f'INDEX(类别_名称,{p})&""')
        hset(ws, f'AN{r}', f'IF(AD{r}="",0,IF(MATCH({esc(f"AD{r}")},类别_名称,0)={p},1,0))')
        hset(ws, f'AE{r}', f'IF(AN{r}=1,INDEX(类别_归类,{p}),"")')
        e = esc(f'AD{r}')
        for col, fld, rng_ in (('AH', '资_收入', inm), ('AI', '资_支出', inm), ('AJ', '资_收入', iny), ('AK', '资_支出', iny)):
            hset(ws, f'{col}{r}', f'IF(AN{r}<>1,0,{R2(f"SUMIFS({fld},资_类别,{e},{rng_})")})')
    counter(ws, 'AO', C0, N_CAT, lambda i: f'$AN{C0 + i}=1')
    sc.set('类别数', f'MIN({cnt("AO", C0, N_CAT)},{N_CAT})')
    CK = dict(月收='AH', 月支='AI', 年收='AJ', 年支='AK')
    for k, col in CK.items():
        fld = '资_收入' if k.endswith('收') else '资_支出'
        rng_ = inm if k.startswith('月') else iny
        sc.set(f'合计_{k}', R2(f'SUMIFS({fld},{rng_},资_归类,"<>内部转账")'))
        sc.set(f'内转_{k}', R2(f'SUMIFS({fld},{rng_},资_归类,"内部转账")'))
        sc.set(f'全部_{k}', R2(f'SUMIFS({fld},{rng_})'))
        sc.set(f'没归类_{k}', R2(f'{sc(f"全部_{k}")}-SUM(${col}${C0}:${col}${C0 + N_CAT - 1})'))
    sc.set('有没归类', 'IF(OR(' + ','.join(f'{sc(f"没归类_{k}")}<>0' for k in CK) + '),1,0)')
    nc, oth = sc('类别数'), sc('有没归类')
    sc.set('显示行数', f'{nc}+2+{oth}')
    sc.set('末行', f'{MO_H3}+{sc("显示行数")}')
    for k in range(1, MO_NC + 1):
        r = MO_C0 + k - 1
        ix = f'$AP{r}'
        hset(ws, f'AP{r}', f'IF({k}<={nc},{kth(k, "AO", C0, N_CAT)},IF({k}={nc}+1,IF({oth}=1,-1,-2),'
                           f'IF({k}={nc}+2,IF({oth}=1,-2,-3),IF(AND({k}={nc}+3,{oth}=1),-3,0))))')
        ws.merge_cells(f'A{r}:B{r}')
        cell(ws, f'A{r}', pick(ix, 'AD', C0, N_CAT, ['"（类别没选或不在清单里的）"', '"合计（不含内部转账）"', '"其中内部转账"']), align=AL)
        cell(ws, f'C{r}', pick(ix, 'AE', C0, N_CAT, ['"？"', '""', '"内部转账"']), align=AC)
        for col, k_ in zip('DEFG', CK):
            cell(ws, f'{col}{r}', pick(ix, CK[k_], C0, N_CAT, [sc(f'没归类_{k_}'), sc(f'合计_{k_}'), sc(f'内转_{k_}')]),
                 fmt=MONEY, align=AR)
        ws.merge_cells(f'H{r}:{LAST}{r}')
        dif = f'ROUND(N(D{r})-N(E{r}),2)'
        cell(ws, f'H{r}', (f'IF({ix}=0,"",IF({ix}=-3,IF(AND({dif}=0,ROUND(N(F{r})-N(G{r}),2)=0),"✓ 转进＝转出",'
                           f'"⚠ 本月转进比转出"&IF({dif}>=0,"多 ","少 ")&TEXT(ABS({dif}),"#,##0.00")'
                           f'&IF(ROUND(N(F{r})-N(G{r}),2)<>0,"；本年差 "&TEXT(N(F{r})-N(G{r}),"#,##0.00;-#,##0.00"),"")&"（只记了一边？）"),'
                           f'IF({ix}=-1,"⚠ 到资金台帐看最右边的校验列",IF({ix}=-2,"＝全部收支－内部转账",'
                           f'IF(C{r}="内部转账","不算在合计里","")))))'), F_NOTE, align=AL)
    c0, c1 = MO_C0, MO_C0 + MO_NC - 1
    z0 = f'$AP{c0}'
    rng = f'A{c0}:{LAST}{c1}'
    cf.add(f'H{c0}:{LAST}{c1}', FormulaRule(formula=[f'LEFT($H{c0},1)="⚠"'], font=Font(bold=True, color='FFC00000')))
    cf.add(rng, FormulaRule(formula=[f'{z0}=-2'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}=-3'], fill=cf_fill('FFDDEBF7'), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}=-1'], font=Font(italic=True, color='FFC00000'), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'AND({z0}>0,$C{c0}="内部转账")'], font=Font(italic=True, color='FF808080'), border=CF_BD,
                            stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}>0'], border=CF_BD))
    hide(ws, *[CL(i) for i in range(CI('K'), CI(mat[-1]) + 1)])
    finish(ws, LAST, f'{MO_H3}:{MO_H3}', sc('末行'), c1)


# ═══════════════════════════════ 费用统计 ═══════════════════════════════
TIP_FEE = ('💡 费用按月统计。黄格 B3 填年份（空着＝截止日那年）。金额按「所属月份」算：资金台帐里填了所属月份的（比如 10 月发 9 月的工资）'
           '算到 9 月，没填按付款月份；只算截止日以前付的。分 6 类：生产成本、销售费用、管理费用、财务费用、税金、其他支出，'
           '每类下面列出【基础资料】里算到这一类的收支类别，段尾小计，最后合计和占比。付供应商货款、买设备、借款还款、内部转账不是费用，不在这里。'
           '右边 ② 选一个类别（空着＝全部费用）、一个月份（空着＝全年）看每一笔。')

F_R0 = 7                                   # ① 第 7 行起
F_NR = len(EXPENSE_KINDS) * (FEE_SEG_MAX + 1) + 1     # 127
F_CAT0 = 7                                 # 隐藏类别表第 7 行起（BA:BD）
F_SEG0 = 3                                 # 隐藏 6 类表（AD:AT）第 3～8 行，第 9 行合计
F_L0 = 7                                   # ② 明细第 7 行起
F_NL = FEE_SHOW + 2
F_K0 = 2                                   # 隐藏：每条资金一行（BX 第几条、BY 符合条件的排序键）
MCOL = [CL(2 + i) for i in range(12)]      # B～M：1～12 月
SUBM = [CL(CI('AI') + i) for i in range(12)]   # 隐藏 6 类 × 12 月：AI～AT


def build_fee(ws):
    LAST, L2 = 'O', 'W'
    widths(ws, dict({'A': 18, 'N': 13.5, 'O': 8, 'P': 2, 'Q': 11.5, 'R': 11, 'S': 13, 'T': 15, 'U': 24, 'V': 13.5, 'W': 9.5},
                    **{c: 12.5 for c in MCOL}))
    title(ws, '=IF(P_公司名称="","",P_公司名称&"　")&"费用统计"', LAST, C_VIEW, TIP_FEE)
    ws.merge_cells(f'Q1:{L2}1')
    put(ws, 'Q1', '② 明细', F_TITLE, fill(C_VIEW), align=AC, border=False)
    ws.merge_cells(f'Q2:{L2}2')
    put(ws, 'Q2', '💡 选一个收支类别（空着＝全部费用）、一个月份（1～12，空着＝全年），列出这一年（按所属月份）的每一笔，按日期排。'
                  '金额＝支出－收入（退回来的钱是负数）。', F_TIP, FILL_TIP, align=ALW, border=False)
    selector(ws, 'A3', '年份', 'B3', None, fmt='0')
    dv = DataValidation(type='whole', operator='between', formula1='2000', formula2='2099', allow_blank=True,
                        showErrorMessage=False, showInputMessage=True, promptTitle='提示', prompt='填年份，比如 2026；空着＝截止日那年')
    ws.add_data_validation(dv)
    dv.add('B3')
    home_link(ws, f'{LAST}3')
    selector(ws, 'Q3', '类别', 'R3', None, '=收支类别列表', prompt='从下拉选收支类别；空着＝全部费用（6 类）')
    ws.merge_cells('R3:S3')
    selector(ws, 'T3', '月份', 'U3', None, '"' + ','.join(str(i) for i in range(1, 13)) + '"', fmt='0',
             prompt='选 1～12 月；空着＝全年')
    ws.row_dimensions[3].height = 26

    keys = ['输入', 'Y', '认出', '截止年月', '段行数', '末行', '超20', '选类别', '类别在', '类别归类', '月输入', '月', '月认出',
            '起月', '止月', '笔数', '显示数', '明细合计']
    sc = Sc(ws, keys)
    v, Y = sc('输入'), sc('Y')
    sc.set('输入', 'IF(ISNUMBER(B3),B3,IFERROR(--TRIM(B3&""),0))')
    sc.set('Y', f'IF(TRIM(B3&"")="",P_年度,IF(AND({v}>=2000,{v}<=2099),INT({v}),IF(AND({v}>=36526,{v}<73051),YEAR({v}),P_年度)))')
    sc.set('认出', f'IF(TRIM(B3&"")="",1,IF(OR(AND({v}>=2000,{v}<=2099),AND({v}>=36526,{v}<73051)),1,0))')
    sc.set('截止年月', 'P_截止年月')
    upto = '资_日期,"<="&P_截止'

    # ── 隐藏 6 类表：AD 类名、AE 类别个数、AF 显示个数、AG 段前行数、AH 全年、AI～AT 1～12 月；第 9 行合计 ──
    S0, S1 = F_SEG0, F_SEG0 + len(EXPENSE_KINDS) - 1
    ST = S1 + 1
    cat_seg = f'$BB${F_CAT0}:$BB${F_CAT0 + N_CAT - 1}'
    for s, kd in enumerate(EXPENSE_KINDS, 1):
        r = S0 + s - 1
        ws[f'AD{r}'] = kd
        ws[f'AD{r}'].font = F_HELP
        hset(ws, f'AE{r}', f'COUNTIF({cat_seg},{s})')
        hset(ws, f'AF{r}', f'MIN(AE{r},{FEE_SEG_MAX})')
        hset(ws, f'AG{r}', '0' if s == 1 else f'AG{r - 1}+AF{r - 1}+1')
        for i, c in enumerate(SUBM):
            hset(ws, f'{c}{r}', R2(f'SUMIFS(资_支出额,资_归类,$AD{r},资_所属年月,{Y}*100+{i + 1},{upto})'))
        hset(ws, f'AH{r}', R2(f'SUM({SUBM[0]}{r}:{SUBM[-1]}{r})'))
    ws[f'AD{ST}'] = '合计'
    ws[f'AD{ST}'].font = F_HELP
    for c in SUBM + ['AH']:
        hset(ws, f'{c}{ST}', R2(f'SUM({c}{S0}:{c}{S1})'))
    KR = f'$AD${S0}:$AD${S1}'
    sc.set('段行数', f'$AG${S1}+$AF${S1}+1')
    sc.set('末行', f'{F_R0 - 1}+{sc("段行数")}+1')
    sc.set('超20', 'IF(OR(' + ','.join(f'$AE${S0 + i}>{FEE_SEG_MAX}' for i in range(len(EXPENSE_KINDS))) + '),1,0)')

    # ── 隐藏类别表（第 F_CAT0 行起，第 i 个类别）：BA 名称、BB 第几类（0＝不是费用/重复）、BC 类内第几个、BD 键＝类×1000＋第几个 ──
    for i in range(N_CAT):
        p, r = i + 1, F_CAT0 + i
        hset(ws, f'BA{r}', f'INDEX(类别_名称,{p})&""')
        hset(ws, f'BB{r}', f'IF(BA{r}="",0,IF(MATCH({esc(f"BA{r}")},类别_名称,0)<>{p},0,IFERROR(MATCH(INDEX(类别_归类,{p}),{KR},0),0)))')
        hset(ws, f'BC{r}', f'IF(BB{r}=0,0,COUNTIF($BB${F_CAT0}:BB{r},BB{r}))')
        hset(ws, f'BD{r}', f'IF(OR(BB{r}=0,BC{r}>{FEE_SEG_MAX}),"",BB{r}*1000+BC{r})')
    KEYR = f'$BD${F_CAT0}:$BD${F_CAT0 + N_CAT - 1}'

    # ── 第 3、4 行 ──
    ws.merge_cells('C3:N3')
    over = '&'.join(f'IF($AE${S0 + i}>{FEE_SEG_MAX},"「{kd}」","")' for i, kd in enumerate(EXPENSE_KINDS))
    put(ws, 'C3', (f'="＝ "&{Y}&" 年"&IF(TRIM(B3&"")="","（空着＝截止日那年）","")'
                   f'&IF({sc("认出")}=0,"（⚠ 没认出来，先按截止日那年）","")'
                   f'&"；按所属月份，只算截止日 "&TEXT(P_截止,"yyyy-mm-dd")&" 以前付的"'
                   f'&IF({sc("超20")}=1,"　⚠ "&{over}&"的类别超过 {FEE_SEG_MAX} 个，只列前 {FEE_SEG_MAX} 个（小计、合计是全部的）","")'),
        F_NOTE, align=ALW, border=False)
    ws.conditional_formatting.add('C3', FormulaRule(formula=['ISNUMBER(FIND("⚠",$C$3))'], font=Font(bold=True, color='FFC00000')))
    kpi(ws, 'A4', '全年费用合计', 'B4', f'$AH${ST}', MONEY0, merge_to='C4')
    for i, kd in enumerate(EXPENSE_KINDS):
        lc, vc = CL(CI('D') + 2 * i), CL(CI('E') + 2 * i)
        kpi(ws, f'{lc}4', kd, f'{vc}4', f'$AH${S0 + i}', MONEY0)
    ws.row_dimensions[4].height = 22

    # ── ① 费用 × 月 ──
    section(ws, 5, 'A', LAST, f'="① "&{Y}&" 年费用（按所属月份；6 类，每类列出收支类别，段尾小计）"', C_VIEW)
    header(ws, 6, [('A', '收支类别')] + [(c, f'{i + 1} 月') for i, c in enumerate(MCOL)] + [('N', '全年'), ('O', '占比')],
           GREEN_H, height=30)
    T = sc('段行数')
    tot_y = f'$AH${ST}'
    for k in range(1, F_NR + 1):
        r = F_R0 + k - 1
        hset(ws, f'BE{r}', k)
        q = f'$BE{r}'
        s, j, t, pos, nm = (f'${c}{r}' for c in ('BF', 'BG', 'BH', 'BI', 'BJ'))
        hset(ws, f'BF{r}', f'IF({q}>{T}+1,0,IF({q}={T}+1,7,MATCH({q}-1,$AG${S0}:$AG${S1},1)))')
        hset(ws, f'BG{r}', f'IF(OR({s}=0,{s}=7),0,{q}-INDEX($AG${S0}:$AG${S1},{s}))')
        hset(ws, f'BH{r}', f'IF({s}=0,0,IF({s}=7,3,IF({j}<=INDEX($AF${S0}:$AF${S1},{s}),1,2)))')
        hset(ws, f'BI{r}', f'IF({t}=1,IFERROR(MATCH({s}*1000+{j},{KEYR},0),0),0)')
        hset(ws, f'BJ{r}', f'IF({t}=1,INDEX(类别_名称,MAX(1,{pos})),IF({t}=2,INDEX({KR},{s}),""))')
        cell(ws, f'A{r}', f'IF({t}=1,{nm},IF({t}=2,"小计："&{nm},IF({t}=3,"费用合计","")))', align=AL)
        for i, c in enumerate(MCOL):
            sub = f'INDEX(${SUBM[i]}${S0}:${SUBM[i]}${S1},{s})'
            f = (f'IF({t}=1,ROUND(SUMIFS(资_支出额,资_类别,{esc(nm)},资_所属年月,{Y}*100+{i + 1},{upto}),2),'
                 f'IF({t}=2,{sub},IF({t}=3,${SUBM[i]}${ST},"")))')
            cell(ws, f'{c}{r}', f, fmt=MONEY0, align=AR)
        cell(ws, f'N{r}', f'IF({t}=0,"",ROUND(SUM(B{r}:M{r}),2))', F_TXTB, MONEY0, AR)
        cell(ws, f'O{r}', f'IF(OR({t}=0,{tot_y}=0),"",N{r}/{tot_y})', F_NOTE, PCT, AR)
    r0, r1 = F_R0, F_R0 + F_NR - 1
    cf = ws.conditional_formatting
    t0 = f'$BH{r0}'
    rng = f'A{r0}:{LAST}{r1}'
    cf.add(rng, FormulaRule(formula=[f'{t0}=3'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=2'], fill=cf_fill('FFDDEBF7'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=1'], border=CF_BD))
    for i, c in enumerate(MCOL):          # 截止日以后的月份：表头灰
        cf.add(f'{c}6', FormulaRule(formula=[f'DATE({Y},{i + 1},1)>P_截止'], fill=cf_fill('FFBFBFBF')))

    # ── ② 明细（右边 Q～W） ──
    sel, mo = sc('选类别'), sc('月')
    sc.set('选类别', 'TRIM(R3&"")')
    sc.set('类别在', f'IF({sel}="",1,IF(COUNTIF(类别_名称,{esc(sel)})>0,1,0))')
    sc.set('类别归类', f'IF({sel}="","",IFERROR(INDEX(类别_归类,MATCH({esc(sel)},类别_名称,0)),""))')
    sc.set('月输入', 'IF(ISNUMBER(U3),U3,IFERROR(--SUBSTITUTE(TRIM(U3&""),"月",""),0))')
    mv = sc('月输入')
    sc.set('月', f'IF(AND({mv}>=1,{mv}<=12,INT({mv})={mv}),{mv},0)')
    sc.set('月认出', f'IF(TRIM(U3&"")="",1,IF({mo}>0,1,0))')
    sc.set('起月', f'{Y}*100+IF({mo}=0,1,{mo})')
    sc.set('止月', f'{Y}*100+IF({mo}=0,12,{mo})')
    K0, K1 = F_K0, F_K0 + N_CASH - 1
    KEYS = f'$BY${K0}:$BY${K1}'
    sc.set('笔数', f'COUNT({KEYS})')
    sc.set('显示数', f'MIN({sc("笔数")},{FEE_SHOW})')
    lo, hi = sc('起月'), sc('止月')
    per = f'资_所属年月,">="&{lo},资_所属年月,"<="&{hi},{upto}'
    all_fee = '+'.join(f'SUMIFS(资_支出额,资_归类,"{kd}",{per})' for kd in EXPENSE_KINDS)
    sc.set('明细合计', f'IF({sel}="",{R2(all_fee)},{R2(f"SUMIFS(资_支出额,资_类别,{esc(sel)},{per})")})')
    for i in range(N_CASH):
        r = K0 + i
        ws[f'BX{r}'] = i + 1
        ws[f'BX{r}'].font = F_HELP
        n = f'$BX{r}'
        x = lambda nm: f'INDEX({nm},{n})'
        ok_cat = f'IF({sel}="",ISNUMBER(MATCH({x("资_归类")},{KR},0)),{x("资_类别")}={sel})'
        hset(ws, f'BY{r}', (f'IF({x("资_排序键")}="","",IF(AND({x("资_所属年月")}>={lo},{x("资_所属年月")}<={hi},'
                            f'{x("资_日期")}<=P_截止,{ok_cat}),{x("资_排序键")},""))'))
    N, NN = sc('笔数'), sc('显示数')
    ws.merge_cells(f'V3:{L2}3')
    put(ws, 'V3', (f'=IF({mo}=0,"全年","")&IF({sc("月认出")}=0,"⚠ 月份填 1～12","")'
                   f'&IF({sc("类别在")}=0,"　⚠ 类别不在清单里","")'), F_NOTE, align=ALW, border=False)
    section(ws, 5, 'Q', L2, (f'="② "&{Y}&" 年"&IF({mo}=0,"全年",{mo}&" 月")&"　"&IF({sel}="","全部费用","「"&{sel}&"」")'
                             f'&"（按所属月份）共 "&{N}&" 笔"&IF({N}>{FEE_SHOW},"，⚠ 只列前 {FEE_SHOW} 笔","")'), C_VIEW)
    ws.merge_cells(f'Q4:{L2}4')
    put(ws, 'Q4', (f'=IF(AND({sel}<>"",{sc("类别归类")}<>"",NOT(ISNUMBER(MATCH({sc("类别归类")},{KR},0)))),'
                   f'"「"&{sel}&"」算到「"&{sc("类别归类")}&"」，不是费用：金额＝支出－收入（负数＝收进来的钱）","")'),
        F_NOTE, align=ALW, border=False)
    header(ws, 6, [('Q', '日期'), ('R', '账户'), ('S', '往来单位'), ('T', '收支类别'), ('U', '摘要'), ('V', '金额'),
                   ('W', '所属月份')], GREEN_H, height=30)
    for k in range(1, F_NL + 1):
        r = F_L0 + k - 1
        ws[f'BK{r}'] = k
        ws[f'BK{r}'].font = F_HELP
        q, y, n, t = f'$BK{r}', f'$BL{r}', f'$BM{r}', f'$BN{r}'
        hset(ws, f'BL{r}', f'IF({q}>{NN},"",SMALL({KEYS},{q}))')
        hset(ws, f'BM{r}', f'IF({y}="",0,MOD({y},{M}))')
        hset(ws, f'BN{r}', f'IF({n}>0,1,IF({q}={NN}+1,3,IF({q}={NN}+2,4,0)))')
        ix = lambda nm: f'INDEX({nm},{n})'
        cell(ws, f'Q{r}', f'IF({t}=1,{ix("资_日期")},IF({t}=3,"合计",IF({t}=4,"共 "&{N}&" 笔","")))', fmt=DATE)
        cell(ws, f'R{r}', f'IF({t}=1,{ix("资_账户")},"")', align=AL)
        cell(ws, f'S{r}', f'IF({t}=1,{ix("资_单位")},"")', align=AL)
        cell(ws, f'T{r}', f'IF({t}=1,{ix("资_类别")},"")', align=AL)
        cell(ws, f'U{r}', f'IF({t}=1,{ix("资_显示摘要")},IF(AND({t}=4,{N}>{FEE_SHOW}),"⚠ 只列了前 {FEE_SHOW} 笔（合计是全部的）",""))',
             align=AL)
        cell(ws, f'V{r}', f'IF({t}=1,{ix("资_支出额")},IF({t}=3,{sc("明细合计")},""))', fmt=MONEY, align=AR)
        ym_ = ix('资_所属年月')
        cell(ws, f'W{r}', f'IF({t}=1,INT({ym_}/100)&"-"&TEXT(MOD({ym_},100),"00"),"")')
    d0, d1 = F_L0, F_L0 + F_NL - 1
    t0 = f'$BN{d0}'
    rng = f'Q{d0}:{L2}{d1}'
    cf.add(rng, FormulaRule(formula=[f'{t0}=3'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=4'], font=Font(color='FF808080'), stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{t0}=1'], border=CF_BD))
    cf.add(f'U{d0}:U{d1}', FormulaRule(formula=[f'AND({t0}=4,LEFT($U{d0},1)="⚠")'], font=Font(bold=True, color='FFC00000')))
    cf.add('Q5', FormulaRule(formula=['ISNUMBER(FIND("⚠",$Q$5))'], font=Font(bold=True, color='FFFFFF00')))
    cf.add('V3', FormulaRule(formula=['ISNUMBER(FIND("⚠",$V$3))'], font=Font(bold=True, color='FFC00000')))
    hide(ws, *[CL(i) for i in range(CI('X'), CI('BY') + 1)])
    finish(ws, LAST, '6:6', sc('末行'), r1, freeze='A7')


def build(wb, ctx=None):
    build_day(wb[SH_DAY])
    build_mon(wb[SH_MON])
    build_fee(wb[SH_FEE])
