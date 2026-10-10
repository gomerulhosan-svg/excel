# -*- coding: utf-8 -*-
"""往来组（橙）：客户往来、供应商往来、对账单，以及隐藏表 _序（对账单的合并排序键）。
   只用 layout 的定义名称、本表格子（对账单还用 _序），不引用录入表。

【客户往来】黄格 C3 起（空＝P_月初）、E3 止（空＝P_截止；晚于截止日按截止日）。
   第 4 行（固定位置，首页可引用）：C4 客户欠款合计（止日欠款只算正数）、E4 预收合计（正数显示）、
   G4 超期没回款家数（位_没回款天数 > P_回款天数）、J4 超期的欠款（这些家到截止日的 位_应收）、M4 待定价发货笔数。
   第 7 行起：位_是客户=1 的每一家，按止日欠款从大到小（预收的排最后，一样多按往来单位表顺序）；后面合计行、「共 N 家」。
   列：序号、客户、期初应收、起日前欠、本期发货、本期回款、止日欠款、预收、已开票、未开票（到截止日）、待定价发货、
   最后发货、最后回款、没回款天数、状态。
【供应商往来】同样：C4 欠供应商合计（正数）、E4 预付合计、G4 待审批合计、J4 已批未付合计。
   列：序号、供应商、期初应付、起日前欠、本期到货、本期付款、止日欠款、预付、已收票、欠票、待定价到货、待审批、已批未付、
   最后到货、最后付款、状态。
   隐藏：AA/AB 标量（第 3 行起）；AD..AL 每家一行（第 i 家在第 6+i 行）：家号、是不是、期初、起日前欠、本期货、本期钱、止日欠、
   分键（止日欠×100 取整）、名次；AM 显示行取第几家（-1 合计、-2 共几家）。
【对账单】黄格 B3 往来单位（下拉；简称、全称、其他叫法都认；空＝第一家客户）、F3 对账方向（客户/供应商；空＝是客户就按客户，否则按供应商）、
   H3 起（空＝P_建账日）、J3 止（空＝P_截止；晚于截止日按截止日）。
   第 6 行抬头（单位不在往来单位表＝大字提示）、第 7 行期间、联系人电话；第 8 行期末大字；第 9 行期初＋本期－本期＝期末；第 10 行本期发票。
   第 11 行表头（每页重复）、第 12 行期初、第 13 行起明细（最多 300 行，同一天先货后钱），紧跟本期合计、共几笔、签字。
   余额＝期初＋SUM(货列 13 行到本行)－SUM(钱列 13 行到本行)。
   隐藏：AA/AB 标量（第 3 行起，期初/本期货/本期钱/期末/待定价/本期票/到止日欠票等）；AD..AH 每行：k、类型、键、货1钱2、第几条。
【_序】A 列（第 2 行起）：第 2～NG+1 行＝货（按方向：销售或采购）第 n 条的键，后面＝资金第 n 条的键；
   键＝日期×SORT_M＋偏移（货的偏移＝n，钱＝NG＋n），所以同一天先货后钱。对账单第 k 行＝SMALL(_序!A, k)。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side
from openpyxl.workbook.defined_name import DefinedName
from common import *
from layout import *          # layout 的 norm / esc 在后面，盖过 common 的
from s_mirror import unit_no

M = SORT_M
NL = 300                                   # 对账单明细最多显示 300 行
NG = max(N_SALE, N_PUR)                    # _序：货占 1..NG，钱占 NG+1..NG+N_CASH
assert NG + N_CASH < SORT_M, '合并排序键的条数要小于 SORT_M'
SEQ = f'{H_SORT}!$A$2:$A${NG + N_CASH + 1}'

_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
LBL = fill('FFD9E1F2')
F_KV = Font(name=YH, sz=10, bold=True, color='FFC00000')
F_BIG = Font(name=YH, sz=16, bold=True, color='FF000000')
F_BIG2 = Font(name=YH, sz=14, bold=True, color='FF1F3864')
F_HEADB = Font(name=YH, sz=11, bold=True, color='FF000000')
DATE_SEL = 'IF(AND(ISNUMBER({x}),{x}>=36526,{x}<73051),INT({x}),{d})'
OK_SEL = 'IF(TRIM({x}&"")="",1,IF(AND(ISNUMBER({x}),{x}>=36526,{x}<73051),1,0))'
QTY = 'General'
PRICE = '0.00##'
DAYS = '0;-0;""'


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def R2(x):
    return f'ROUND({x},2)'


class Sc:
    """隐藏标量：AA 列名字、AB 列值（第 r0 行起）"""

    def __init__(self, ws, keys, r0=3, cl='AA', cv='AB'):
        self.ws, self.cv = ws, cv
        self.r = {k: r0 + i for i, k in enumerate(keys)}
        for k, r in self.r.items():
            ws[f'{cl}{r}'] = k
            ws[f'{cl}{r}'].font = F_HELP

    def __call__(self, k):
        return f'${self.cv}${self.r[k]}'

    def ext(self, k):
        """别的表引用这个标量"""
        return f'{self.ws.title}!${self.cv}${self.r[k]}'

    def set(self, k, f):
        c = self.ws[f'{self.cv}{self.r[k]}']
        c.value = '=' + f
        c.font = F_HELP


def hset(ws, coord, f):
    ws[coord] = f if isinstance(f, (int, float)) else '=' + f
    ws[coord].font = F_HELP


def cell(ws, coord, f, font=F_TXT, fmt=None, align=AC):
    c = ws[coord]
    c.value = '=' + f
    c.font = font
    if fmt:
        c.number_format = fmt
    c.alignment = align
    return c


def kpi(ws, lbl_cell, lbl, val_cell, f, fmt=MONEY):
    put(ws, lbl_cell, lbl, F_KPI_L, LBL, align=ACW)
    put(ws, val_cell, '=' + f, F_KV, FILL_AUTO, fmt, AR)


def finish(ws, last, hdr_rows, start, end_cell, cap_row, freeze, landscape):
    ws.freeze_panes = freeze
    print_setup(ws, hdr_rows, landscape=landscape)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A${start}:INDEX({q}!${last}$1:${last}${cap_row},{q}!{end_cell})')


# ═══════════════════════════════ 客户往来 / 供应商往来 ═══════════════════════════════
P_H, P_L0 = 6, 7                       # 表头行、清单第一行
P_N = N_UNIT + 2                       # 每家一行＋合计＋共几家
U0 = P_L0                              # 隐藏每家表：第 i 家在第 U0+i-1 行
HU = dict(家='AD', 是='AE', 期初='AF', 前欠='AG', 货='AH', 钱='AI', 止欠='AJ', 分='AK', 名次='AL', 取='AM')


def _rng(col):
    return f'${col}${U0}:${col}${U0 + N_UNIT - 1}'


TIP_AR = ('💡 每家客户欠多少、这段时间发了多少货、回了多少款。黄格选起止日期：起空着＝截止日那个月 1 号，止空着＝截止日'
          '（截止日在首页改；晚于截止日的按截止日算），右边灰字是实际用的日期。只列【往来单位】里类型是客户（或客户和供应商）的，'
          '按止日欠款从大到小排，预收（欠款是负数）的排最后。期初应收是建账那天的数；起日前欠＝期初＋起日以前的发货－回款；'
          '止日欠款＝起日前欠＋本期发货－本期回款。已开票、未开票、待定价、最后发货、最后回款、没回款天数都是到截止日的数，'
          '没回款天数超过【基础资料】里设的天数标红。要给哪家对账，去【对账单】选它。')
TIP_AP = ('💡 欠每家供应商多少、这段时间到了多少货、付了多少款。黄格选起止日期：起空着＝截止日那个月 1 号，止空着＝截止日'
          '（晚于截止日的按截止日算），右边灰字是实际用的日期。只列【往来单位】里类型是供应商（或客户和供应商）的，'
          '按止日欠款从大到小排，预付（欠款是负数）的排最后。起日前欠＝期初应付＋起日以前的到货－付款；止日欠款＝起日前欠＋本期到货－本期付款。'
          '已收票、欠票、待定价、待审批（付款审批里老板还没批的申请金额）、已批未付（批了还没付完的）、最后到货、最后付款都是到截止日的数。')


def build_party(ws, cust):
    who = '客户' if cust else '供应商'
    LAST = 'O' if cust else 'P'
    if cust:
        W = dict(A=5, B=16, C=13.5, D=13.5, E=13.5, F=13.5, G=13.5, H=6, I=13.5, J=13.5, K=8, L=12.5, M=12.5, N=8, O=24)
    else:
        W = dict(A=5, B=16, C=13.5, D=13.5, E=13.5, F=13.5, G=13.5, H=6, I=13.5, J=13.5, K=8, L=13.5, M=13.5, N=12.5,
                 O=12.5, P=22)
    widths(ws, W)
    title(ws, f'=IF(P_公司名称="","",P_公司名称&"　")&"{who}往来"', LAST, C_WL, TIP_AR if cust else TIP_AP)
    selector(ws, 'B3', '起日期', 'C3', None, fmt=DATE)
    selector(ws, 'D3', '止日期', 'E3', None, fmt=DATE)
    dv_date(ws, 'C3')
    dv_date(ws, 'E3')
    home_link(ws, f'{LAST}3')
    ws.row_dimensions[3].height = 30

    keys = ['起', '起认出', '止输入', '止', '止认出', '家数', '末行', '回款天数',
            't_期初', 't_前欠', 't_货', 't_钱', 't_止欠', 't_票', 't_欠票', 't_待定价', 'k_欠', 'k_预']
    keys += ['k_超家', 'k_超额'] if cust else ['t_待审批', 't_已批未付']
    sc = Sc(ws, keys)
    S0, S1 = sc('起'), sc('止')
    sc.set('起', DATE_SEL.format(x='C3', d='P_月初'))
    sc.set('起认出', OK_SEL.format(x='C3'))
    sc.set('止输入', DATE_SEL.format(x='E3', d='P_截止'))
    sc.set('止', f'MIN({sc("止输入")},P_截止)')
    sc.set('止认出', OK_SEL.format(x='E3'))
    sc.set('回款天数', 'P_回款天数')

    # 口径
    flag = '位_是客户' if cust else '位_是供应商'
    open_ = '位_期初应收' if cust else '位_期初应付'
    if cust:
        goods = lambda u, cond: f'SUMIFS(销_金额,销_单位号,{u},销_计入往来,1,{cond})'
        money = lambda u, cond: f'SUMIFS(资_净额,资_单位号,{u},资_归类,"收客户货款",{cond})'
        gd, md = '销_日期', '资_日期'
    else:
        goods = lambda u, cond: f'SUMIFS(采_金额,采_单位号,{u},采_计入往来,1,{cond})'
        money = lambda u, cond: f'SUMIFS(资_支出额,资_单位号,{u},资_归类,"付供应商货款",{cond})'
        gd, md = '采_日期', '资_日期'

    # ── 隐藏：每家一行 ──
    H = HU
    for i in range(N_UNIT):
        u, r = i + 1, U0 + i
        c = lambda k: f'${H[k]}{r}'
        ws[f'{H["家"]}{r}'] = u
        ws[f'{H["家"]}{r}'].font = F_HELP
        hset(ws, f'{H["是"]}{r}', f'IF(INDEX({flag},{u})=1,1,0)')
        ok = f'{c("是")}=1'
        hset(ws, f'{H["期初"]}{r}', f'IF({ok},INDEX({open_},{u}),0)')
        lt = f'"<"&{S0}'
        bef = R2(c('期初') + '+' + goods(u, f'{gd},{lt}') + '-' + money(u, f'{md},{lt}'))
        hset(ws, f'{H["前欠"]}{r}', f'IF({ok},{bef},0)')
        hset(ws, f'{H["货"]}{r}', f'IF({ok},{R2(goods(u, dr(gd, S0, S1)))},0)')
        hset(ws, f'{H["钱"]}{r}', f'IF({ok},{R2(money(u, dr(md, S0, S1)))},0)')
        hset(ws, f'{H["止欠"]}{r}', R2(f'{c("前欠")}+{c("货")}-{c("钱")}'))
        hset(ws, f'{H["分"]}{r}', f'IF({ok},ROUND({c("止欠")}*100,0),"")')
        kk = c('分')
        above = (f'+COUNTIF(${H["分"]}${U0}:{H["分"]}{r - 1},{kk})' if i else '')
        hset(ws, f'{H["名次"]}{r}', f'IF({ok},1+COUNTIF({_rng(H["分"])},">"&{kk}){above},"")')
    rg = lambda k: _rng(H[k])
    sc.set('家数', f'COUNTIF({rg("是")},1)')
    sc.set('末行', f'{P_L0 - 1}+{sc("家数")}+2')
    for k, hk in (('t_期初', '期初'), ('t_前欠', '前欠'), ('t_货', '货'), ('t_钱', '钱'), ('t_止欠', '止欠')):
        sc.set(k, R2(f'SUMIFS({rg(hk)},{rg("是")},1)'))
    sc.set('k_欠', R2(f'SUMIFS({rg("止欠")},{rg("是")},1,{rg("止欠")},">0")'))
    sc.set('k_预', R2(f'-SUMIFS({rg("止欠")},{rg("是")},1,{rg("止欠")},"<0")'))
    if cust:
        sc.set('t_票', R2(f'SUMIFS(位_已开票,{flag},1)'))
        sc.set('t_欠票', R2(f'SUMIFS(位_未开票,{flag},1)'))
        sc.set('t_待定价', f'SUMIFS(位_待定价发货,{flag},1)')
        sc.set('k_超家', f'COUNTIFS({flag},1,位_没回款天数,">"&P_回款天数)')
        sc.set('k_超额', R2(f'SUMIFS(位_应收,{flag},1,位_没回款天数,">"&P_回款天数)'))
    else:
        sc.set('t_票', R2(f'SUMIFS(位_已收票,{flag},1)'))
        sc.set('t_欠票', R2(f'SUMIFS(位_欠票,{flag},1)'))
        sc.set('t_待定价', f'SUMIFS(位_待定价到货,{flag},1)')
        sc.set('t_待审批', R2(f'SUMIFS(位_待审批,{flag},1)'))
        sc.set('t_已批未付', R2(f'SUMIFS(位_已批未付,{flag},1)'))

    # ── 第 3、4、5 行 ──
    nl = 'N' if cust else 'O'
    ws.merge_cells(f'F3:{nl}3')
    put(ws, 'F3', (f'="实际用的："&TEXT({S0},"yyyy-mm-dd")&" ～ "&TEXT({S1},"yyyy-mm-dd")'
                   f'&IF(TRIM(C3&"")="","（起空着＝截止日那个月 1 号）","")&IF(TRIM(E3&"")="","（止空着＝截止日）","")'
                   f'&IF({sc("起认出")}=0,"　⚠ 起日期没认出，按截止日那个月 1 号","")'
                   f'&IF({sc("止认出")}=0,"　⚠ 止日期没认出，按截止日","")'
                   f'&IF({sc("止输入")}>P_截止,"　⚠ 止晚于截止日 "&TEXT(P_截止,"yyyy-mm-dd")&"，按截止日算","")'
                   f'&IF({S0}>{S1},"　✗ 起晚于止：本期都是 0，请改日期","")'), F_NOTE, align=ALW, border=False)
    ws.conditional_formatting.add('F3', FormulaRule(formula=['OR(ISNUMBER(FIND("⚠",$F$3)),ISNUMBER(FIND("✗",$F$3)))'],
                                                    font=Font(bold=True, color='FFC00000')))
    if cust:
        kpi(ws, 'B4', '客户欠款合计\n（只算欠的）', 'C4', sc('k_欠'))
        kpi(ws, 'D4', '预收合计', 'E4', sc('k_预'))
        put(ws, 'F4', '="超 "&P_回款天数&" 天没回款"', F_KPI_L, LBL, align=ACW)
        put(ws, 'G4', '=' + sc('k_超家'), F_KV, FILL_AUTO, '0" 家"', AR)
        kpi(ws, 'I4', '超期的欠款\n（到截止日）', 'J4', sc('k_超额'))
        ws.merge_cells('K4:L4')
        put(ws, 'K4', '待定价发货', F_KPI_L, LBL, align=ACW)
        ws['L4'].border = BD
        put(ws, 'M4', '=' + sc('t_待定价'), F_KV, FILL_AUTO, '0" 笔"', AR)
    else:
        kpi(ws, 'B4', '欠供应商合计\n（只算欠的）', 'C4', sc('k_欠'))
        kpi(ws, 'D4', '预付合计', 'E4', sc('k_预'))
        kpi(ws, 'F4', '待审批合计', 'G4', sc('t_待审批'))
        kpi(ws, 'I4', '已批未付合计', 'J4', sc('t_已批未付'))
        ws.merge_cells('K4:L4')
        put(ws, 'K4', '待定价到货', F_KPI_L, LBL, align=ACW)
        ws['L4'].border = BD
        put(ws, 'M4', '=' + sc('t_待定价'), F_KV, FILL_AUTO, '0" 笔"', AR)
    ws.row_dimensions[4].height = 32
    ws.merge_cells(f'A5:{LAST}5')
    if cust:
        t5 = ('="按止日欠款从大到小排，预收（负数）的排最后。已开票、未开票、待定价、最后发货、最后回款、没回款天数都是到截止日 "'
              '&TEXT(P_截止,"yyyy-mm-dd")&" 的数；没回款天数超过 "&P_回款天数&" 天标红（天数在【基础资料】改）。"')
    else:
        t5 = ('="按止日欠款从大到小排，预付（负数）的排最后。已收票、欠票、待定价、待审批、已批未付、最后到货、最后付款都是到截止日 "'
              '&TEXT(P_截止,"yyyy-mm-dd")&" 的数。"')
    put(ws, 'A5', t5, F_NOTE, align=ALW, border=False)
    ws.row_dimensions[5].height = 20

    # ── 表头 ──
    if cust:
        heads = [('A', '序号'), ('B', '客户'), ('C', '期初应收\n（建账日）'), ('D', '起日前欠'), ('E', '本期发货'), ('F', '本期回款'),
                 ('G', '止日欠款\n（负数＝预收）'), ('H', '预收'), ('I', '已开票'), ('J', '未开票'), ('K', '待定价\n发货(笔)'),
                 ('L', '最后发货'), ('M', '最后回款'), ('N', '没回款\n天数'), ('O', '状态')]
    else:
        heads = [('A', '序号'), ('B', '供应商'), ('C', '期初应付\n（建账日）'), ('D', '起日前欠'), ('E', '本期到货'), ('F', '本期付款'),
                 ('G', '止日欠款\n（负数＝预付）'), ('H', '预付'), ('I', '已收票'), ('J', '欠票'), ('K', '待定价\n到货(笔)'),
                 ('L', '待审批'), ('M', '已批未付'), ('N', '最后到货'), ('O', '最后付款'), ('P', '状态')]
    header(ws, P_H, heads, C_WL, height=36)

    # ── 清单 ──
    n = sc('家数')
    for k in range(1, P_N + 1):
        r = P_L0 + k - 1
        ix = f'${H["取"]}{r}'
        hset(ws, f'{H["取"]}{r}', f'IF({k}<={n},IFERROR(MATCH({k},{rg("名次")},0),0),IF({k}={n}+1,-1,IF({k}={n}+2,-2,0)))')
        pk = lambda rng_, tot: f'IF({ix}=0,"",IF({ix}>0,INDEX({rng_},{ix}),IF({ix}=-1,{tot},"")))'
        pos = lambda rng_: f'IF({ix}>0,IF(INDEX({rng_},{ix})>0,INDEX({rng_},{ix}),""),"")'
        cell(ws, f'A{r}', f'IF({ix}>0,{k},"")', F_NOTE)
        cell(ws, f'B{r}', f'IF({ix}>0,INDEX(位_简称,{ix}),IF({ix}=-1,"合计",IF({ix}=-2,"共 "&{n}&" 家{who}","")))', align=AL)
        for col, hk, tk in (('C', '期初', 't_期初'), ('D', '前欠', 't_前欠'), ('E', '货', 't_货'), ('F', '钱', 't_钱'),
                            ('G', '止欠', 't_止欠')):
            cell(ws, f'{col}{r}', pk(rg(hk), sc(tk)), fmt=MONEY, align=AR)
        zq = f'INDEX({rg("止欠")},{ix})'
        cell(ws, f'H{r}', f'IF({ix}>0,IF({zq}<-0.005,"{"预收" if cust else "预付"}",""),"")', F_TXTB)
        if cust:
            cell(ws, f'I{r}', pk('位_已开票', sc('t_票')), fmt=MONEY, align=AR)
            cell(ws, f'J{r}', pk('位_未开票', sc('t_欠票')), fmt=MONEY, align=AR)
            cell(ws, f'K{r}', pk('位_待定价发货', sc('t_待定价')), fmt=INT)
            cell(ws, f'L{r}', pos('位_最后发货'), fmt=DATE)
            cell(ws, f'M{r}', pos('位_最后回款'), fmt=DATE)
            cell(ws, f'N{r}', f'IF({ix}>0,INDEX(位_没回款天数,{ix}),"")', fmt=DAYS)
            st = (f'IF({zq}<-0.005,"、预收","")&IF(N(INDEX(位_没回款天数,{ix}))>P_回款天数,"、超 "&P_回款天数&" 天没回款","")'
                  f'&IF(INDEX(位_待定价发货,{ix})>0,"、有待定价","")')
            cell(ws, f'O{r}', f'IF({ix}>0,MID({st},2,100),"")', align=AL)
        else:
            cell(ws, f'I{r}', pk('位_已收票', sc('t_票')), fmt=MONEY, align=AR)
            cell(ws, f'J{r}', pk('位_欠票', sc('t_欠票')), fmt=MONEY, align=AR)
            cell(ws, f'K{r}', pk('位_待定价到货', sc('t_待定价')), fmt=INT)
            cell(ws, f'L{r}', pk('位_待审批', sc('t_待审批')), fmt=MONEY, align=AR)
            cell(ws, f'M{r}', pk('位_已批未付', sc('t_已批未付')), fmt=MONEY, align=AR)
            cell(ws, f'N{r}', pos('位_最后到货'), fmt=DATE)
            cell(ws, f'O{r}', pos('位_最后付款'), fmt=DATE)
            st = (f'IF({zq}<-0.005,"、预付","")&IF(INDEX(位_待定价到货,{ix})>0,"、有待定价","")'
                  f'&IF(INDEX(位_待审批,{ix})>0,"、有待审批","")')
            cell(ws, f'P{r}', f'IF({ix}>0,MID({st},2,100),"")', align=AL)
    a0, a1 = P_L0, P_L0 + P_N - 1
    z0 = f'${H["取"]}{a0}'
    cf = ws.conditional_formatting
    rng = f'A{a0}:{LAST}{a1}'
    cf.add(rng, FormulaRule(formula=[f'{z0}=-1'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}=-2'], font=Font(color='FF808080'), stopIfTrue=True))
    if cust:
        cf.add(f'N{a0}:N{a1}', FormulaRule(formula=[f'AND({z0}>0,N($N{a0})>{sc("回款天数")})'],
                                         font=Font(bold=True, color='FFC00000'), border=CF_BD, stopIfTrue=True))
    sc_ = 'O' if cust else 'P'
    cf.add(f'{sc_}{a0}:{sc_}{a1}', FormulaRule(formula=[f'AND({z0}>0,ISNUMBER(FIND("超",${sc_}{a0})))'],
                                             font=Font(bold=True, color='FFC00000'), border=CF_BD, stopIfTrue=True))
    cf.add(f'{sc_}{a0}:{sc_}{a1}', FormulaRule(formula=[f'AND({z0}>0,${sc_}{a0}<>"")'],
                                             font=Font(bold=True, color='FFC65911'), border=CF_BD, stopIfTrue=True))
    cf.add(f'H{a0}:H{a1}', FormulaRule(formula=[f'AND({z0}>0,$H{a0}<>"")'], font=Font(bold=True, color='FF2F75B5'),
                                     fill=cf_fill('FFDDEBF7'), border=CF_BD, stopIfTrue=True))
    cf.add(rng, FormulaRule(formula=[f'{z0}>0'], border=CF_BD))
    hide(ws, *[CL(i) for i in range(CI('AA'), CI(H['取']) + 1)])
    finish(ws, LAST, f'{P_H}:{P_H}', 1, sc('末行'), a1, 'C7', True)


# ═══════════════════════════════ 对账单 ═══════════════════════════════
ST_LAST = 'J'
ST_H, ST_O, ST_D0 = 11, 12, 13          # 表头、期初行、明细第一行
ST_N = NL + 7                           # 明细 300 行＋本期合计、共几笔、空、签字、空、空、日期
ST_D1 = ST_D0 + ST_N - 1
HS = dict(k='AD', 类型='AE', 键='AF', 源='AG', 条='AH')

TIP_ST = ('💡 给客户、供应商对账，A4 竖着打（一页宽，表头每页重复），打出来请对方核对盖章。黄格：往来单位（下拉选，也可以打简称、全称、'
          '其他叫法；空着＝第一家客户）、对账方向（客户/供应商；空着＝按单位类型，是客户就按客户对，不然按供应商）、起（空＝建账日）、'
          '止（空＝截止日；晚于截止日的按截止日）。明细按日期排，同一天先列货、后列钱；余额＝期初＋发货－回款逐笔往下算。'
          '待定价的货（还没定单价）金额按 0 算，备注写「单价待定」。客户：货＝销售登记，钱＝资金台帐里「收客户货款」（退给客户的是负数）；'
          '供应商：货＝采购登记，钱＝「付供应商货款」。打印只印第 6 行往下。')


def build_stmt(ws):
    LAST = ST_LAST
    W = dict(A=12.5, B=13, C=20, D=8, E=6, F=10, G=13.5, H=13.5, I=14, J=20, K=10)
    widths(ws, W)
    title(ws, '对账单', LAST, C_WL, TIP_ST)
    home_link(ws, 'K1')
    selector(ws, 'A3', '往来单位', 'B3', None, '=往来单位列表',
             prompt='从下拉选；也可以打简称、全称、其他叫法；空着＝第一家客户')
    ws.merge_cells('B3:C3')
    selector(ws, 'D3', '对账方向', 'F3', None, '"客户,供应商"', prompt='空着＝按单位类型（是客户就按客户，不然按供应商）')
    ws.merge_cells('D3:E3')
    ws['E3'].border = BD
    selector(ws, 'G3', '起', 'H3', None, fmt=DATE)
    selector(ws, 'I3', '止', 'J3', None, fmt=DATE)
    dv_date(ws, 'H3')
    dv_date(ws, 'J3')
    ws.row_dimensions[3].height = 26

    keys = ['名', '名规', 'u找', '首客户', 'u', '是客户', '是供应商', '方向输入', 'dir', '起', '起认出', '止输入', '止', '止认出',
            'lo', 'hi', '期初', '货', '钱', '期末', '待定价', '票', '欠票', '笔数', 'n', '末行', '简称', '抬头名', '联系人', '电话']
    sc = Sc(ws, keys)
    U, D, S0, S1 = sc('u'), sc('dir'), sc('起'), sc('止')
    sc.set('名', 'TRIM(B3&"")')
    sc.set('名规', norm('B3'))
    sc.set('u找', unit_no(sc('名规')))
    sc.set('首客户', 'IFERROR(MATCH(1,位_是客户,0),0)')
    sc.set('u', f'IF({sc("名")}="",{sc("首客户")},{sc("u找")})')
    sc.set('是客户', f'IF({U}=0,0,INDEX(位_是客户,{U}))')
    sc.set('是供应商', f'IF({U}=0,0,INDEX(位_是供应商,{U}))')
    sc.set('方向输入', 'TRIM(F3&"")')
    sc.set('dir', f'IF({sc("方向输入")}="客户",1,IF({sc("方向输入")}="供应商",2,IF({sc("是客户")}=1,1,2)))')
    sc.set('起', DATE_SEL.format(x='H3', d='P_建账日'))
    sc.set('起认出', OK_SEL.format(x='H3'))
    sc.set('止输入', DATE_SEL.format(x='J3', d='P_截止'))
    sc.set('止', f'MIN({sc("止输入")},P_截止)')
    sc.set('止认出', OK_SEL.format(x='J3'))
    sc.set('lo', f'{S0}*{M}')
    sc.set('hi', f'({S1}+1)*{M}')
    bef = lambda nm: f'{nm},"<"&{S0}'
    per = lambda nm: dr(nm, S0, S1)
    upto = lambda nm: f'{nm},"<="&{S1}'
    cs = dict(
        期初=(f'IF({sc("是客户")}=1,INDEX(位_期初应收,{U}),0)+SUMIFS(销_金额,销_单位号,{U},销_计入往来,1,{bef("销_日期")})'
              f'-SUMIFS(资_净额,资_单位号,{U},资_归类,"收客户货款",{bef("资_日期")})',
              f'IF({sc("是供应商")}=1,INDEX(位_期初应付,{U}),0)+SUMIFS(采_金额,采_单位号,{U},采_计入往来,1,{bef("采_日期")})'
              f'-SUMIFS(资_支出额,资_单位号,{U},资_归类,"付供应商货款",{bef("资_日期")})'),
        货=(f'SUMIFS(销_金额,销_单位号,{U},销_计入往来,1,{per("销_日期")})',
            f'SUMIFS(采_金额,采_单位号,{U},采_计入往来,1,{per("采_日期")})'),
        钱=(f'SUMIFS(资_净额,资_单位号,{U},资_归类,"收客户货款",{per("资_日期")})',
            f'SUMIFS(资_支出额,资_单位号,{U},资_归类,"付供应商货款",{per("资_日期")})'),
        票=(f'SUMIFS(票_价税合计,票_单位号,{U},票_方向,"开出",票_计入往来,1,{per("票_日期")})',
            f'SUMIFS(票_价税合计,票_单位号,{U},票_方向,"收到",票_计入往来,1,{per("票_日期")})'),
        欠票=(f'INDEX(位_期初未开票,{U})+SUMIFS(销_金额,销_单位号,{U},销_计入往来,1,销_开票,1,{upto("销_日期")})'
              f'-SUMIFS(票_价税合计,票_单位号,{U},票_方向,"开出",票_计入往来,1,{upto("票_日期")})',
              f'INDEX(位_期初欠票,{U})+SUMIFS(采_金额,采_单位号,{U},采_计入往来,1,采_开票,1,{upto("采_日期")})'
              f'-SUMIFS(票_价税合计,票_单位号,{U},票_方向,"收到",票_计入往来,1,{upto("票_日期")})'),
    )
    for k, (a, b) in cs.items():
        sc.set(k, f'IF({U}=0,0,IF({D}=1,{R2(a)},{R2(b)}))')
    sc.set('期末', R2(f'{sc("期初")}+{sc("货")}-{sc("钱")}'))
    sc.set('待定价', (f'IF({U}=0,0,IF({D}=1,COUNTIFS(销_单位号,{U},销_计入往来,1,销_待定价,1,{per("销_日期")}),'
                    f'COUNTIFS(采_单位号,{U},采_计入往来,1,采_待定价,1,{per("采_日期")})))'))
    sc.set('笔数', f'COUNT({SEQ})')
    sc.set('n', f'MIN({sc("笔数")},{NL})')
    sc.set('末行', f'{ST_D0 + 6}+{sc("n")}')
    sc.set('简称', f'IF({U}=0,{sc("名")},INDEX(位_简称,{U}))')
    sc.set('抬头名', f'IF({U}=0,"",IF(INDEX(位_全称,{U})="",INDEX(位_简称,{U}),INDEX(位_全称,{U})))')
    sc.set('联系人', f'IF({U}=0,"",INDEX(位_联系人,{U})&"")')
    sc.set('电话', f'IF({U}=0,"",INDEX(位_电话,{U})&"")')

    # ── 第 4、5 行：实际用的、提醒 ──
    put(ws, 'A4', '实际用的', F_NOTE, align=AC, border=False)
    ws.merge_cells('B4:C4')
    put(ws, 'B4', (f'=IF({U}=0,IF({sc("名")}="","⚠ 往来单位表里还没有客户","⚠ 不在【往来单位】表里"),'
                   f'"✓ "&{sc("简称")}&"（"&INDEX(位_类型,{U})&"）"&IF({sc("名")}="","　空着＝第一家客户",""))'),
        F_AUTOB, FILL_AUTO, align=ACW)
    ws.merge_cells('D4:F4')
    put(ws, 'D4', f'=IF({D}=1,"按客户对","按供应商对")&IF({sc("方向输入")}="","（按单位类型）","")', F_AUTOB, FILL_AUTO, align=AC)
    ws['E4'].border = ws['F4'].border = BD
    put(ws, 'H4', f'={S0}', F_AUTOB, FILL_AUTO, DATE, AC)
    put(ws, 'J4', f'={S1}', F_AUTOB, FILL_AUTO, DATE, AC)
    ws.row_dimensions[4].height = 30
    ws.merge_cells(f'A5:{LAST}5')
    warn = (f'IF({sc("起认出")}=0,"⚠ 起日期没认出，按建账日。","")&IF({sc("止认出")}=0,"⚠ 止日期没认出，按截止日。","")'
            f'&IF({sc("止输入")}>P_截止,"⚠ 止晚于截止日 "&TEXT(P_截止,"yyyy-mm-dd")&"，按截止日算。","")'
            f'&IF({S0}>{S1},"✗ 起晚于止，请改日期。","")'
            f'&IF({sc("笔数")}>{NL},"⚠ 共 "&{sc("笔数")}&" 笔，只列了前 {NL} 笔（合计是全部的），请缩短日期分几张对。","")')
    put(ws, 'A5', (f'=IF({warn}="","起空着＝建账日 "&TEXT(P_建账日,"yyyy-mm-dd")&"，止空着＝截止日 "&TEXT(P_截止,"yyyy-mm-dd")'
                   f'&"；往来单位空着＝第一家客户。下面第 6 行起是打印出来的对账单。",{warn})'), F_NOTE, align=ALW, border=False)
    ws.conditional_formatting.add('A5', FormulaRule(formula=['OR(LEFT($A$5,1)="⚠",LEFT($A$5,1)="✗")'],
                                                    font=Font(bold=True, color='FFC00000')))
    ws.row_dimensions[5].height = 22

    # ── 抬头 ──
    ws.merge_cells(f'A6:{LAST}6')
    put(ws, 'A6', (f'=IF({U}=0,IF({sc("名")}="","⚠【往来单位】表里还没有客户，先去加一家",'
                   f'"⚠ 「"&{sc("名")}&"」不在【往来单位】表里，先去【往来单位】加一行（简称、全称、其他叫法都认）"),'
                   f'IF(P_公司名称="","",P_公司名称&" 与 ")&{sc("抬头名")}&" 对账单")'), F_BIG, align=ACW, border=False)
    ws.conditional_formatting.add('A6', FormulaRule(formula=['LEFT($A$6,1)="⚠"'], font=Font(bold=True, color='FFC00000')))
    ws.row_dimensions[6].height = 40
    ws.merge_cells('A7:E7')
    put(ws, 'A7', f'="对账期间："&TEXT({S0},"yyyy-mm-dd")&" 至 "&TEXT({S1},"yyyy-mm-dd")', F_HEADB, align=AL, border=False)
    ws.merge_cells(f'F7:{LAST}7')
    put(ws, 'F7', f'="对方联系人："&{sc("联系人")}&"　　电话："&{sc("电话")}', F_TXT, align=AR, border=False)
    ws.row_dimensions[7].height = 22
    E = sc('期末')
    amt = lambda x: f'TEXT({x},"#,##0.00")'
    big = (f'IF({U}=0,"","截至 "&TEXT({S1},"yyyy-mm-dd")&"，"&IF({D}=1,'
           f'IF({E}>0.005,"贵公司尚欠我公司 "&{amt(E)}&" 元",IF({E}<-0.005,"我公司预收贵公司 "&{amt("-" + E)}&" 元","双方两清")),'
           f'IF({E}>0.005,"我公司尚欠贵公司 "&{amt(E)}&" 元",IF({E}<-0.005,"我公司预付贵公司 "&{amt("-" + E)}&" 元","双方两清"))))')
    ws.merge_cells(f'A8:{LAST}8')
    put(ws, 'A8', '=' + big, F_BIG2, fill('FFFFF2CC'), align=ACW, border=False)
    ws.row_dimensions[8].height = 32
    ws.merge_cells(f'A9:{LAST}9')
    put(ws, 'A9', (f'=IF({U}=0,"","期初余额 "&{amt(sc("期初"))}&" ＋ 本期"&IF({D}=1,"发货","到货")&" "&{amt(sc("货"))}'
                   f'&" － 本期"&IF({D}=1,"回款","付款")&" "&{amt(sc("钱"))}&" ＝ 期末余额 "&{amt(E)}&" 元"'
                   f'&IF({sc("待定价")}>0,"；另有 "&{sc("待定价")}&" 笔"&IF({D}=1,"发货","到货")&"单价待定（金额没算）",""))'),
        F_TXT, align=ALW, border=False)
    ws.row_dimensions[9].height = 20
    ws.merge_cells(f'A10:{LAST}10')
    q = sc('欠票')
    put(ws, 'A10', (f'=IF({U}=0,"",IF({D}=1,"本期我公司开出发票 "&{amt(sc("票"))}&" 元；截至 "&TEXT({S1},"yyyy-mm-dd")'
                    f'&IF({q}<-0.005,"，发票多开了 "&{amt("-" + q)}&" 元","，还有 "&{amt(q)}&" 元没开票"),'
                    f'"本期收到贵公司发票 "&{amt(sc("票"))}&" 元；截至 "&TEXT({S1},"yyyy-mm-dd")'
                    f'&IF({q}<-0.005,"，贵公司发票多开了 "&{amt("-" + q)}&" 元","，贵公司还欠我公司发票 "&{amt(q)}&" 元")))'),
        F_TXT, align=ALW, border=False)
    ws.row_dimensions[10].height = 20

    # ── 表头、期初行 ──
    header(ws, ST_H, [('A', '日期'), ('B', '单号'), ('C', '品名规格'), ('D', '数量'), ('E', '单位'), ('F', '单价'),
                      ('G', f'=IF({D}=1,"发货金额","到货金额")'), ('H', f'=IF({D}=1,"回款金额","付款金额")'), ('I', '余额'),
                      ('J', '备注')], C_WL, height=30)
    r = ST_O
    for col in 'ABCDEFGHIJ':
        put(ws, f'{col}{r}', None, F_TXTB, fill('FFF2F2F2'), align=AC)
    put(ws, f'A{r}', f'=IF({U}=0,"","期初余额")', F_TXTB, fill('FFF2F2F2'), align=AL)
    put(ws, f'C{r}', f'=IF({U}=0,"","截至 "&TEXT({S0}-1,"yyyy-mm-dd"))', F_TXTB, fill('FFF2F2F2'), align=AL)
    put(ws, f'I{r}', f'=IF({U}=0,"",{sc("期初")})', F_TXTB, fill('FFF2F2F2'), MONEY, AR)

    # ── 明细（活动区：明细 → 本期合计 → 共几笔 → 签字） ──
    n, N_ = sc('n'), sc('笔数')
    for kk in range(ST_N):
        k, r = kk + 1, ST_D0 + kk
        t, key, s, i = (f'${HS[x]}{r}' for x in ('类型', '键', '源', '条'))
        ws[f'{HS["k"]}{r}'] = k
        ws[f'{HS["k"]}{r}'].font = F_HELP
        hset(ws, f'{HS["类型"]}{r}', f'IF({k}<={n},1,IF({k}={n}+1,2,IF({k}={n}+2,3,IF({k}={n}+4,4,IF({k}={n}+7,5,0)))))')
        hset(ws, f'{HS["键"]}{r}', f'IF({t}=1,SMALL({SEQ},{k}),0)')
        hset(ws, f'{HS["源"]}{r}', f'IF({t}<>1,0,IF(MOD({key},{M})<={NG},1,2))')
        hset(ws, f'{HS["条"]}{r}', f'IF({s}=1,MOD({key},{M}),IF({s}=2,MOD({key},{M})-{NG},0))')
        g = lambda a, b: f'IF({D}=1,INDEX({a},{i}),INDEX({b},{i}))'
        z = lambda nm: f'INDEX({nm},{i})'
        no_s = (f'IF({z("销_出库单号")}="",{z("销_合同号")},IF({z("销_合同号")}="",{z("销_出库单号")},'
                f'{z("销_合同号")}&"/"&{z("销_出库单号")}))')
        cell(ws, f'A{r}', f'IF({t}=1,INT({key}/{M}),IF({t}=2,"本期合计",IF({t}=3,"共 "&{N_}&" 笔","")))', fmt=DATE)
        cell(ws, f'B{r}', (f'IF({s}=1,IF({D}=1,{no_s},{z("采_采购单号")}),IF({s}=2,IF({z("资_关联键")}="","",'
                           f'MID({z("资_关联键")},2,100)),""))'), align=AL)
        cash_lbl = (f'IF({D}=1,IF({z("资_净额")}<0,"退款","回款"),IF({z("资_支出额")}<0,"退回","付款"))')
        cell(ws, f'C{r}', (f'IF({s}=1,{g("销_品名规格", "采_品名规格")},IF({s}=2,{cash_lbl},IF({t}=3,'
                           f'IF({N_}>{NL},"⚠ 只列了前 {NL} 笔",""),IF({t}=4,"我方（盖章）：",IF({t}=5,"日期：","")))))'), align=AL)
        cell(ws, f'D{r}', f'IF({s}=1,IF({g("销_数量", "采_数量")}=0,"",{g("销_数量", "采_数量")}),"")', fmt=QTY, align=AR)
        cell(ws, f'E{r}', f'IF({s}=1,{g("销_计量单位", "采_计量单位")},"")')
        cell(ws, f'F{r}', f'IF({s}=1,IF({g("销_单价", "采_单价")}=0,"",{g("销_单价", "采_单价")}),"")', fmt=PRICE, align=AR)
        cell(ws, f'G{r}', f'IF({s}=1,{g("销_金额", "采_金额")},IF({t}=2,{sc("货")},""))', fmt=MONEY, align=AR)
        cell(ws, f'H{r}', f'IF({s}=2,{g("资_净额", "资_支出额")},IF({t}=2,{sc("钱")},""))', fmt=MONEY, align=AR)
        cell(ws, f'I{r}', (f'IF({t}=1,ROUND({sc("期初")}+SUM(G${ST_D0}:G{r})-SUM(H${ST_D0}:H{r}),2),'
                           f'IF({t}=2,{E},""))'), fmt=MONEY, align=AR)
        bz = g('销_备注', '采_备注')
        cell(ws, f'J{r}', (f'IF({s}=1,IF({g("销_待定价", "采_待定价")}=1,"单价待定"&IF({bz}="","","；"&{bz}),{bz}),'
                           f'IF({s}=2,{z("资_显示摘要")},IF({t}=2,IF({sc("待定价")}>0,"另有 "&{sc("待定价")}&" 笔单价待定",""),'
                           f'IF({t}=4,"对方确认（盖章）：",IF({t}=5,"日期：","")))))'), align=AL)
    reg = f'A{ST_D0}:{LAST}{ST_D1}'
    t0 = f'${HS["类型"]}{ST_D0}'
    cf = ws.conditional_formatting
    cf.add(reg, FormulaRule(formula=[f'{t0}=2'], fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD, stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'{t0}=3'], font=Font(color='FF808080'), stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'OR({t0}=4,{t0}=5)'], font=Font(bold=True), stopIfTrue=True))
    cf.add(f'J{ST_D0}:J{ST_D1}', FormulaRule(formula=[f'AND({t0}=1,LEFT($J{ST_D0},4)="单价待定")'],
                                             font=Font(bold=True, color='FFC00000'), border=CF_BD, stopIfTrue=True))
    cf.add(reg, FormulaRule(formula=[f'{t0}=1'], border=CF_BD))
    for rr in range(ST_D0, ST_D1 + 1):
        ws.row_dimensions[rr].height = 18
    hide(ws, *[CL(i) for i in range(CI('AA'), CI(HS['条']) + 1)])
    finish(ws, LAST, f'{ST_H}:{ST_H}', 6, sc('末行'), ST_D1, 'A5', False)
    ws.print_options.horizontalCentered = True
    return sc


def build_seq(wb, sc):
    """_序 A 列：货（按方向）第 n 条在第 n+1 行，资金第 n 条在第 NG+n+1 行"""
    ws = wb[H_SORT]
    ws['A1'] = f'对账单排序键：日期×{M}＋偏移（货 1～{NG}，钱 {NG}＋n）'
    ws['B1'] = '第 2～' + str(NG + 1) + ' 行＝货，往下＝资金'
    u, d, lo, hi = sc.ext('u'), sc.ext('dir'), sc.ext('lo'), sc.ext('hi')
    rng = lambda k: f'IF(AND(N({k})>={lo},N({k})<{hi}),{k},"")'
    for n in range(1, NG + 1):
        sale = (f'IF(INDEX(销_单位号,{n})={u},{rng(f"INDEX(销_排序键,{n})")},"")' if n <= N_SALE else '""')
        pur = (f'IF(INDEX(采_单位号,{n})={u},{rng(f"INDEX(采_排序键,{n})")},"")' if n <= N_PUR else '""')
        ws[f'A{n + 1}'] = f'=IF({u}=0,"",IF({d}=1,{sale},{pur}))'
    for n in range(1, N_CASH + 1):
        k = f'INDEX(资_排序键,{n})'
        ws[f'A{NG + n + 1}'] = (f'=IF({u}=0,"",IF(AND(INDEX(资_单位号,{n})={u},INDEX(资_归类,{n})=IF({d}=1,"收客户货款","付供应商货款"),'
                                f'N({k})>={lo},N({k})<{hi}),{k}+{NG},""))')


def build(wb, ctx=None):
    build_party(wb[SH_AR], True)
    build_party(wb[SH_AP], False)
    sc = build_stmt(wb[SH_STMT])
    build_seq(wb, sc)
