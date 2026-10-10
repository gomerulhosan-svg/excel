# -*- coding: utf-8 -*-
"""录入表（蓝）：资金台帐、销售登记、采购登记、发票登记、付款审批、采购计划；基础（灰）：基础资料、往来单位、固定支出。
   录入表只放录入列，外加最右边几列只给人看的显示列（别的公式不引用它们，插进来的行没有这几格也不影响汇总）。"""
import datetime as dt
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from layout import *
from common import *


def _d(v):
    if v in (None, ''):
        return None
    if isinstance(v, (dt.date, dt.datetime)):
        return dt.datetime(v.year, v.month, v.day)
    try:
        return dt.datetime.strptime(str(v)[:10], '%Y-%m-%d')
    except ValueError:
        return v


def dv_dec(ws, sqref, what='金额', lo='-999999999', hi='999999999', extra=''):
    dv = DataValidation(type='decimal', operator='between', formula1=lo, formula2=hi, allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what,
                        error=f'{what}要填数字（不要带文字、单位）{extra}')
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_day(ws, sqref):
    """日期：只收真日期（输 10-8 或 2026/10/8；10.8 会被当成小数，拦下来）"""
    dv = DataValidation(type='date', operator='between', formula1='36526', formula2='73050', allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle='日期',
                        error='输 10-8 或 2026/10/8 就行（不要用 10.8，那是小数）')
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_len(ws, sqref, n=50):
    dv = DataValidation(type='textLength', operator='lessThanOrEqual', formula1=str(n), allow_blank=True, showErrorMessage=True,
                        errorStyle='stop', errorTitle='太长', error=f'不要超过 {n} 个字')
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_whole(ws, sqref, lo, hi, what):
    dv = DataValidation(type='whole', operator='between', formula1=str(lo), formula2=str(hi), allow_blank=True,
                        showErrorMessage=True, errorStyle='stop', errorTitle=what, error=f'{what}填 {lo}～{hi} 的整数')
    ws.add_data_validation(dv)
    dv.add(sqref)


def dv_fixed(ws, sqref, items, prompt=None, stop=True):
    return dv_list(ws, sqref, '"' + ','.join(items) + '"', prompt=prompt, stop=stop)


FILL_EX = fill('FFEDEDED')
F_LBL = fill('FFD9E1F2')
H_AUTO = 'FF8EA9DB'        # 自动列表头
H_BOSS = 'FF7030A0'        # 老板填的列表头
F_WARN = Font(name=YH, sz=10, bold=True, color='FFC65911')


def _example_grey(ws, rows, r0, last_in):
    """示例行（备注以「示例」开头）灰底，正式用时整行删掉"""
    for i, x in enumerate(rows):
        if str(x.get('备注') or '').startswith('示例'):
            for c in range(1, CI(last_in) + 1):
                ws.cell(row=r0 + i, column=c).fill = FILL_EX


def _frame(ws, last_col, heads, color, tip, W, nrows, fm, al, freeze, auto=(), title_col=None, hdr_h=44):
    widths(ws, W)
    title(ws, ws.title, title_col or last_col, color, tip)
    header(ws, HDR, heads, color, height=hdr_h)
    style_rows(ws, R0, R0 + nrows - 1, [c for c, _ in heads], auto=auto, fmts=fm, aligns=al)
    ws.freeze_panes = freeze
    ws.auto_filter.ref = f'A{HDR}:{last_col}{R0 + nrows - 1}'
    for col in auto:
        ws[f'{col}{HDR}'].fill = fill(H_AUTO)


def _write_rows(ws, rows, cols, date_keys=('日期',), text_keys=()):
    for i, x in enumerate(rows):
        r = R0 + i
        for k, col in cols.items():
            v = x.get(k)
            if v in (None, ''):
                continue
            if k in date_keys:
                v = _d(v)
            if k in text_keys:
                v = str(v)
            ws[f'{col}{r}'] = v


def _print_to_last(ws, last_col, par_key, cap):
    """打印区域只到最后一行有内容的（_参 里算的行号），不打几千行空表格线"""
    q = f"'{ws.title}'"
    cell = PAR[par_key]
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${last_col}$1:${last_col}${R0 + cap - 1},{H_PAR}!${cell[0]}${cell[1:]})')


def _kpi(ws, row, cells):
    """第 3 行：标签（合并两格）＋ 值"""
    for c, lbl, f, fmt in cells:
        put(ws, f'{c}{row}', lbl, F_KPI_L, F_LBL, align=AC)
        ws.merge_cells(f'{c}{row}:{CL(CI(c) + 1)}{row}')
        put(ws, f'{CL(CI(c) + 2)}{row}', f, F_AUTOB, FILL_AUTO, fmt, AR)


def _show(ws, nshow, cols_formulas, fmts, aligns=None):
    """显示列：第 n 条＝ROW()-ROW(表头)；超出容量或没内容显示空"""
    n = f'ROW()-ROW($A${HDR})'
    aligns = aligns or {}
    for i in range(nshow):
        r = R0 + i
        for col, f in cols_formulas.items():
            c = ws[f'{col}{r}']
            c.value = '=IFERROR(' + f.format(n=n) + ',"")'
            c.fill = FILL_AUTO
            c.font = F_AUTO
            if col in fmts:
                c.number_format = fmts[col]
            if col in aligns:
                c.alignment = aligns[col]


def _check_cf(ws, col, last):
    rng = f'{col}{R0}:{col}{last}'
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{R0},1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT({col}{R0},1)="⚠"'], font=F_WARN))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{col}{R0}="✓"'], font=Font(name=YH, sz=10, color='FF70AD47')))


def _nofx(ws, key_col, chk_col, cell, last, merge_to=None):
    """「右边没公式的行」提示：有日期、问题列却是空的（插进来的行）"""
    f = (f'=IF(COUNTIFS({key_col}{R0}:{key_col}{last},"<>",{chk_col}{R0}:{chk_col}{last},"")>0,'
         f'COUNTIFS({key_col}{R0}:{key_col}{last},"<>",{chk_col}{R0}:{chk_col}{last},"")&" 行右边没公式：选上一行右边灰色几格往下拉","")')
    put(ws, cell, f, F_RED, align=AL, border=False)
    if merge_to:
        ws.merge_cells(f'{cell}:{merge_to}')


def _check(prefix):
    return f'IF(INDEX({prefix}_末行,{{n}})=0,"",IF(INDEX({prefix}_校验,{{n}})="","✓",INDEX({prefix}_校验,{{n}})))'


def _star(k, req):
    return f'{k}★' if k in req else k


# ───────────────────────── 资金台帐 ─────────────────────────
def build_cash(ws, ctx):
    C, S = CASH_COLS, CASH_SHOW
    heads = [(C['日期'], '日期★'), (C['账户'], '账户★'), (C['类别'], '收支类别★'), (C['往来单位'], '往来单位\n（收付货款必选）'),
             (C['摘要'], '摘要'), (C['收入'], '收入'), (C['支出'], '支出'),
             (C['关联单号'], '关联单号\n（付审批过的款、固定支出选）'), (C['所属月份'], '所属月份\n（工资电费填几月；空＝付款月）'),
             (C['备注'], '备注'), (S['余额'], '本账户余额\n（自动，按上下顺序）'), (S['单位类型'], '单位类型\n（自动）'),
             (S['归类'], '算到哪\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 7 个账户混在一张表录，一笔一行，收入、支出只填一个。收客户货款、付供应商货款一定要选往来单位（冲欠款）；'
           '付审批过的款，关联单号选那张审批单。内部转账记两行：转出账户记支出、转进账户记收入，类别都选「内部转账」。'
           '右边自动出这个账户的即时余额（按表里上下顺序算，按日期往下录就对了）。可以按账户、往来单位、收支类别、单位类型筛选，第 3 行是筛出来的合计。灰底是示例行。')
    W = {C['日期']: 11, C['账户']: 10, C['类别']: 14, C['往来单位']: 14, C['摘要']: 24, C['收入']: 12, C['支出']: 12,
         C['关联单号']: 16, C['所属月份']: 12, C['备注']: 16, S['余额']: 14, S['单位类型']: 9, S['归类']: 11, S['校验']: 40}
    fm = {C['日期']: DATE, C['收入']: MONEY, C['支出']: MONEY, C['关联单号']: '@', C['所属月份']: '[<=12]0"月";[>=36526]yyyy-mm;0'}
    al = {C['往来单位']: AL, C['摘要']: AL, C['备注']: AL, C['收入']: AR, C['支出']: AR, C['关联单号']: AL}
    rows = ctx.get('cash', [])
    _frame(ws, CASH_LAST, heads, C_IN, tip, W, N_CASH, fm, al, f'C{R0}', auto=tuple(S.values()), title_col='J')
    last = R0 + N_CASH + 2000
    put(ws, 'A3', '筛选后看得见的行合计', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells('A3:E3')
    for col in (C['收入'], C['支出']):
        put(ws, f'{col}3', f'=SUBTOTAL(9,{col}{HDR}:{col}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["关联单号"]}3', '收入－支出', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["所属月份"]}3', f'={C["收入"]}3-{C["支出"]}3', F_AUTOB, FILL_AUTO, MONEY, AR)
    _nofx(ws, 'A', S['校验'], f'{C["备注"]}3', last, f'{S["校验"]}3')
    _write_rows(ws, rows, C, text_keys=('关联单号',))
    nshow = N_CASH
    _show(ws, nshow, {
        S['余额']: 'INDEX(资_余额,{n})',
        S['单位类型']: 'INDEX(资_单位类型,{n})&""',
        S['归类']: 'IF(INDEX(资_末行,{n})=0,"",INDEX(资_归类,{n}))',
        S['校验']: _check('资'),
    }, {S['余额']: MONEY}, {S['校验']: AL, S['余额']: AR})
    lastshow = R0 + nshow - 1
    _check_cf(ws, S['校验'], lastshow)
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_CASH - 1}'
    dv_day(ws, rng(C['日期']))
    dv_list(ws, rng(C['账户']), '=账户列表', prompt='从【基础资料】②账户里选')
    dv_list(ws, rng(C['类别']), '=收支类别列表', prompt='从【基础资料】③收支类别里选；收付货款选「收客户货款」「付供应商货款」')
    dv_list(ws, rng(C['往来单位']), '=往来单位列表', stop=False,
            prompt='收付货款要选【往来单位】里的客户、供应商；房东、电费户名等可以直接打')
    dv_list(ws, rng(C['关联单号']), '=关联单号列表', stop=False,
            prompt='付审批过的款选审批单（只列批了没付完的）；付固定支出可以选编号（不选也行）')
    for k in ('收入', '支出'):
        dv_dec(ws, rng(C[k]), extra='；10000+200 要写成 =10000+200')
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(CASH_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, CASH_LAST, '资金末行', N_CASH)


# ───────────────────────── 销售登记、采购登记 ─────────────────────────
def build_sale(ws, ctx):
    C, S = SALE_COLS, SALE_SHOW
    heads = [(C['日期'], '发货日期★'), (C['客户'], '客户★'), (C['合同号'], '合同/订单号'), (C['产品'], '产品名称'),
             (C['规格'], '规格型号'), (C['数量'], '数量'), (C['计量单位'], '单位'), (C['单价'], '单价'),
             (C['金额'], '金额\n（空＝数量×单价）'), (C['业务类型'], '业务类型\n（空＝自产加工）'), (C['开票'], '开票\n（空＝按单位设置）'),
             (C['出库单号'], '出库单号'), (C['业务员'], '业务员'), (C['备注'], '备注'),
             (S['计算金额'], '金额\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 发一次货登一行（快工单一张出库单有几个产品就登几行），客户从下拉选。还没定价的，单价、金额先空着，定了价回到原行补上（发货日期不要改）。'
           '买成品回来卖的，业务类型选「外购成品」。退货、折让、质量扣款记负数行（产品名称写原因）。灰底是示例行。')
    W = {C['日期']: 11, C['客户']: 13, C['合同号']: 14, C['产品']: 18, C['规格']: 14, C['数量']: 9, C['计量单位']: 6,
         C['单价']: 9, C['金额']: 12, C['业务类型']: 10, C['开票']: 8, C['出库单号']: 15, C['业务员']: 8, C['备注']: 16,
         S['计算金额']: 12, S['校验']: 38}
    fm = {C['日期']: DATE, C['数量']: '#,##0.##', C['单价']: '#,##0.00##', C['金额']: MONEY, C['合同号']: '@', C['出库单号']: '@'}
    al = {C['客户']: AL, C['产品']: AL, C['规格']: AL, C['备注']: AL, C['合同号']: AL, C['出库单号']: AL}
    for k in ('数量', '单价', '金额'):
        al[C[k]] = AR
    rows = ctx.get('sales', [])
    _frame(ws, SALE_LAST, heads, C_IN, tip, W, N_SALE, fm, al, f'C{R0}', auto=tuple(S.values()), title_col='N')
    last = R0 + N_SALE + 2000
    put(ws, 'A3', '筛选后看得见的行合计', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells('A3:E3')
    put(ws, f'{C["数量"]}3', f'=SUBTOTAL(9,{C["数量"]}{HDR}:{C["数量"]}{last})', F_AUTOB, FILL_AUTO, '#,##0.##', AR)
    put(ws, f'{C["单价"]}3', '金额', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["金额"]}3', f'=SUBTOTAL(9,{S["计算金额"]}{HDR}:{S["计算金额"]}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["业务类型"]}3', '待定价（全部）', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["开票"]}3', '=COUNTIFS(销_待定价,1,销_有效,1)&" 笔"', F_AUTOB, FILL_AUTO, align=AR)
    _nofx(ws, 'A', S['校验'], f'{C["出库单号"]}3', last, f'{S["校验"]}3')
    _write_rows(ws, rows, C, text_keys=('合同号', '出库单号'))
    _show(ws, N_SALE, {
        S['计算金额']: 'IF(INDEX(销_末行,{n})=0,"",IF(INDEX(销_待定价,{n})=1,"待定价",INDEX(销_金额,{n})))',
        S['校验']: _check('销'),
    }, {S['计算金额']: MONEY}, {S['计算金额']: AR, S['校验']: AL})
    lastshow = R0 + N_SALE - 1
    _check_cf(ws, S['校验'], lastshow)
    ws.conditional_formatting.add(f'{S["计算金额"]}{R0}:{S["计算金额"]}{lastshow}',
                                  FormulaRule(formula=[f'{S["计算金额"]}{R0}="待定价"'], font=F_WARN))
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_SALE - 1}'
    dv_day(ws, rng(C['日期']))
    dv_list(ws, rng(C['客户']), '=客户列表', prompt='从【往来单位】的客户里选；新客户先去【往来单位】加一行')
    for k in ('数量', '单价', '金额'):
        dv_dec(ws, rng(C[k]), k)
    dv_fixed(ws, rng(C['业务类型']), BIZ_TYPES, prompt='空着＝自产加工；买成品回来卖的选外购成品')
    dv_fixed(ws, rng(C['开票']), YESNO, prompt='这一行要不要开票；空着＝按【往来单位】里这家的设置')
    dv_list(ws, rng(C['业务员']), '=人员列表', stop=False)
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(SALE_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, SALE_LAST, '销售末行', N_SALE)


def build_pur(ws, ctx):
    C, S = PUR_COLS, PUR_SHOW
    heads = [(C['日期'], '到货日期★'), (C['供应商'], '供应商★'), (C['采购单号'], '采购单号'), (C['物料'], '物料名称'),
             (C['规格'], '规格'), (C['数量'], '数量'), (C['计量单位'], '单位'), (C['单价'], '单价'),
             (C['金额'], '金额\n（空＝数量×单价）'), (C['采购类别'], '采购类别\n（空＝按供应商默认）'), (C['开票'], '要票\n（空＝按单位设置）'),
             (C['备注'], '备注'), (S['计算金额'], '金额\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 到一次货登一行（快工单入库单没有金额，按送货单、磅单、结算单填）。外协、钢材月底才定价的，先填数量，单价、金额空着，定了价回到原行补。'
           '当场付清、没在【往来单位】建档的零星料钱不用登这里，资金台帐记「零星料钱（现买）」就行。退货、扣款记负数。灰底是示例行。')
    W = {C['日期']: 11, C['供应商']: 13, C['采购单号']: 14, C['物料']: 18, C['规格']: 14, C['数量']: 9, C['计量单位']: 6,
         C['单价']: 9, C['金额']: 12, C['采购类别']: 11, C['开票']: 8, C['备注']: 22, S['计算金额']: 12, S['校验']: 38}
    fm = {C['日期']: DATE, C['数量']: '#,##0.###', C['单价']: '#,##0.00##', C['金额']: MONEY, C['采购单号']: '@'}
    al = {C['供应商']: AL, C['物料']: AL, C['规格']: AL, C['备注']: AL, C['采购单号']: AL}
    for k in ('数量', '单价', '金额'):
        al[C[k]] = AR
    rows = ctx.get('purchases', [])
    _frame(ws, PUR_LAST, heads, C_IN, tip, W, N_PUR, fm, al, f'C{R0}', auto=tuple(S.values()), title_col='L')
    last = R0 + N_PUR + 2000
    put(ws, 'A3', '筛选后看得见的行合计', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells('A3:E3')
    put(ws, f'{C["数量"]}3', f'=SUBTOTAL(9,{C["数量"]}{HDR}:{C["数量"]}{last})', F_AUTOB, FILL_AUTO, '#,##0.###', AR)
    put(ws, f'{C["单价"]}3', '金额', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["金额"]}3', f'=SUBTOTAL(9,{S["计算金额"]}{HDR}:{S["计算金额"]}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["采购类别"]}3', '待定价（全部）', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["开票"]}3', '=COUNTIFS(采_待定价,1,采_有效,1)&" 笔"', F_AUTOB, FILL_AUTO, align=AR)
    _nofx(ws, 'A', S['校验'], f'{C["备注"]}3', last, f'{S["校验"]}3')
    _write_rows(ws, rows, C, text_keys=('采购单号',))
    _show(ws, N_PUR, {
        S['计算金额']: 'IF(INDEX(采_末行,{n})=0,"",IF(INDEX(采_待定价,{n})=1,"待定价",INDEX(采_金额,{n})))',
        S['校验']: _check('采'),
    }, {S['计算金额']: MONEY}, {S['计算金额']: AR, S['校验']: AL})
    lastshow = R0 + N_PUR - 1
    _check_cf(ws, S['校验'], lastshow)
    ws.conditional_formatting.add(f'{S["计算金额"]}{R0}:{S["计算金额"]}{lastshow}',
                                  FormulaRule(formula=[f'{S["计算金额"]}{R0}="待定价"'], font=F_WARN))
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_PUR - 1}'
    dv_day(ws, rng(C['日期']))
    dv_list(ws, rng(C['供应商']), '=供应商列表', prompt='从【往来单位】的供应商里选；新供应商先去【往来单位】加一行')
    for k in ('数量', '单价', '金额'):
        dv_dec(ws, rng(C[k]), k)
    dv_fixed(ws, rng(C['采购类别']), PUR_TYPES, prompt='空着＝【往来单位】里这家的默认采购类别（再空＝原材料）')
    dv_fixed(ws, rng(C['开票']), YESNO, prompt='这一行供应商要不要给票；空着＝按【往来单位】里这家的设置')
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(PUR_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, PUR_LAST, '采购末行', N_PUR)


# ───────────────────────── 发票登记 ─────────────────────────
def build_inv(ws, ctx):
    C, S = INV_COLS, INV_SHOW
    heads = [(C['日期'], '开票日期★'), (C['方向'], '开出/收到\n（空＝按单位类型）'), (C['往来单位'], '往来单位★'),
             (C['发票号码'], '发票号码'), (C['发票类型'], '专票/普票'), (C['税率'], '税率\n（空＝13%）'), (C['价税合计'], '价税合计★'),
             (C['税额'], '税额\n（空＝自动算）'), (C['对应单号'], '对应合同/单号'), (C['备注'], '备注'),
             (S['不含税'], '不含税金额\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 只登本公司真实开出、收到的发票，一张一行（不用税局导出）。开出＝我们开给客户，收到＝供应商开给我们；'
           '单位只是客户或只是供应商的，开出/收到可以不填。红字发票价税合计填负数。发票号码这一列是文本格式，直接输。灰底是示例行。')
    W = {C['日期']: 11, C['方向']: 11, C['往来单位']: 14, C['发票号码']: 23, C['发票类型']: 9, C['税率']: 8, C['价税合计']: 13,
         C['税额']: 12, C['对应单号']: 16, C['备注']: 18, S['不含税']: 13, S['校验']: 38}
    fm = {C['日期']: DATE, C['价税合计']: MONEY, C['税额']: MONEY, C['发票号码']: '@', C['对应单号']: '@', C['税率']: '0%'}
    al = {C['往来单位']: AL, C['发票号码']: AL, C['对应单号']: AL, C['备注']: AL, C['价税合计']: AR, C['税额']: AR}
    rows = ctx.get('invoices', [])
    _frame(ws, INV_LAST, heads, C_IN, tip, W, N_INV, fm, al, f'D{R0}', auto=tuple(S.values()), title_col='J')
    last = R0 + N_INV + 2000
    put(ws, 'A3', '筛选后看得见的行合计', F_KPI_L, F_LBL, align=AC)
    ws.merge_cells('A3:F3')
    put(ws, f'{C["价税合计"]}3', f'=SUBTOTAL(9,{C["价税合计"]}{HDR}:{C["价税合计"]}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    put(ws, f'{C["税额"]}3', '不含税', F_KPI_L, F_LBL, align=AC)
    put(ws, f'{C["对应单号"]}3', f'=SUBTOTAL(9,{S["不含税"]}{HDR}:{S["不含税"]}{last})', F_AUTOB, FILL_AUTO, MONEY, AR)
    _nofx(ws, 'A', S['校验'], f'{C["备注"]}3', last, f'{S["校验"]}3')
    for x in rows:                                   # 税率按小数写进去（显示 13%）
        t = x.get('税率')
        if isinstance(t, str) and t.endswith('%'):
            x['税率'] = float(t[:-1]) / 100
    _write_rows(ws, rows, C, text_keys=('发票号码', '对应单号'))
    _show(ws, N_INV, {
        S['不含税']: 'IF(INDEX(票_末行,{n})=0,"",INDEX(票_不含税,{n}))',
        S['校验']: _check('票'),
    }, {S['不含税']: MONEY}, {S['不含税']: AR, S['校验']: AL})
    _check_cf(ws, S['校验'], R0 + N_INV - 1)
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_INV - 1}'
    dv_day(ws, rng(C['日期']))
    dv_fixed(ws, rng(C['方向']), INV_DIRS, prompt='开出＝开给客户；收到＝供应商开来的。空着＝按往来单位类型')
    dv_list(ws, rng(C['往来单位']), '=往来单位列表', prompt='从【往来单位】里选')
    dv_fixed(ws, rng(C['发票类型']), INV_KINDS)
    dv_fixed(ws, rng(C['税率']), TAX_RATES, stop=False, prompt='空着＝13%；也可以直接输 0.13 或 13')
    for k in ('价税合计', '税额'):
        dv_dec(ws, rng(C[k]), k)
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(INV_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, INV_LAST, '发票末行', N_INV)


# ───────────────────────── 付款审批 ─────────────────────────
def build_apv(ws, ctx):
    C, S = APV_COLS, APV_SHOW
    heads = [(C['单号'], '单号★'), (C['申请日期'], '申请日期★'), (C['申请人'], '申请人'), (C['收款单位'], '收款单位★'),
             (C['付款内容'], '付款内容'), (C['申请金额'], '申请金额★'), (C['计划付款日期'], '计划付款日期\n（空＝申请日期）'),
             (C['付款账户'], '付款账户'), (C['审批结果'], '审批结果\n（老板选）'), (C['批准金额'], '批准金额\n（老板；空＝申请的数）'),
             (C['审批意见'], '审批意见\n（老板）'), (C['审批日期'], '审批日期\n（老板）'), (C['备注'], '备注'),
             (S['现在欠款'], '这家现在欠\n（自动）'), (S['欠票'], '欠票\n（自动）'), (S['已付'], '已付\n（自动）'),
             (S['还差'], '批了还没付\n（自动）'), (S['状态'], '状态\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 要付的款都在这里申请一行：货款、外协款，还有每个月计划要付的电费、税、借款利息等（计划付款日期定在哪个月）。'
           '老板在紫色几列选审批结果（同意 / 部分同意 / 不同意 / 暂缓），部分同意填批准金额。'
           '付款时在【资金台帐】那一笔的「关联单号」选这张单，已付、状态自动出。【待审批付款单】可以打印给老板签字。第 3 行有下一个单号。灰底是示例行。')
    W = {C['单号']: 12, C['申请日期']: 11, C['申请人']: 8, C['收款单位']: 14, C['付款内容']: 20, C['申请金额']: 12,
         C['计划付款日期']: 12, C['付款账户']: 10, C['审批结果']: 10, C['批准金额']: 12, C['审批意见']: 16, C['审批日期']: 11,
         C['备注']: 14, S['现在欠款']: 13, S['欠票']: 12, S['已付']: 12, S['还差']: 12, S['状态']: 9, S['校验']: 36}
    fm = {C['申请日期']: DATE, C['计划付款日期']: DATE, C['审批日期']: DATE, C['申请金额']: MONEY, C['批准金额']: MONEY, C['单号']: '@'}
    al = {C['收款单位']: AL, C['付款内容']: AL, C['审批意见']: AL, C['备注']: AL, C['申请金额']: AR, C['批准金额']: AR}
    rows = ctx.get('approvals', [])
    _frame(ws, APV_LAST, heads, C_IN, tip, W, N_APV, fm, al, f'E{R0}', auto=tuple(S.values()), title_col='M')
    for k in APV_BOSS:
        ws[f'{C[k]}{HDR}'].fill = fill(H_BOSS)
        for r in range(R0, R0 + N_APV):
            ws[f'{C[k]}{r}'].fill = fill('FFF3EAF9')
    _kpi(ws, 3, [('A', '下一个单号', '=P_下一审批号', '@'),
                 ('D', '待审批合计', '=SUMIFS(批_申请金额,批_状态码,0)', MONEY),
                 ('G', '批了还没付', '=SUM(批_未付)', MONEY)])
    _nofx(ws, 'B', S['校验'], f'{C["备注"]}3', R0 + N_APV + 2000, f'{S["校验"]}3')
    _write_rows(ws, rows, C, date_keys=('申请日期', '计划付款日期', '审批日期'), text_keys=('单号',))
    code = 'INDEX(批_状态码,{n})'
    _show(ws, N_APV, {
        S['现在欠款']: 'IF(INDEX(批_末行,{n})=0,"",INDEX(批_现在欠款,{n}))',
        S['欠票']: 'IF(INDEX(批_末行,{n})=0,"",INDEX(批_欠票,{n}))',
        S['已付']: 'IF(INDEX(批_末行,{n})=0,"",INDEX(批_已付,{n}))',
        S['还差']: f'IF(OR({code}=1,{code}=2),INDEX(批_未付,{{n}}),"")',
        S['状态']: 'INDEX(批_状态,{n})&""',
        S['校验']: _check('批'),
    }, {S['现在欠款']: MONEY, S['欠票']: MONEY, S['已付']: MONEY, S['还差']: MONEY},
        {S['校验']: AL, S['现在欠款']: AR, S['欠票']: AR, S['已付']: AR, S['还差']: AR})
    lastshow = R0 + N_APV - 1
    _check_cf(ws, S['校验'], lastshow)
    st = S['状态']
    rng_st = f'{st}{R0}:{st}{lastshow}'
    ws.conditional_formatting.add(rng_st, FormulaRule(formula=[f'{st}{R0}="待审批"'], font=F_WARN, fill=fill('FFFFF2CC')))
    ws.conditional_formatting.add(rng_st, FormulaRule(formula=[f'OR({st}{R0}="待付款",{st}{R0}="部分已付")'],
                                                      font=Font(name=YH, sz=10, bold=True, color='FF2F75B5')))
    ws.conditional_formatting.add(f'A{R0}:{C["备注"]}{lastshow}',
                                  FormulaRule(formula=[f'OR(${st}{R0}="不同意",${st}{R0}="已付")'], font=Font(name=YH, sz=10, color='FF808080')))
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_APV - 1}'
    for k in ('申请日期', '计划付款日期', '审批日期'):
        dv_day(ws, rng(C[k]))
    dv_list(ws, rng(C['申请人']), '=人员列表', stop=False)
    dv_list(ws, rng(C['收款单位']), '=往来单位列表', stop=False, prompt='供应商从下拉选；电费、税务局等可以直接打')
    for k in ('申请金额', '批准金额'):
        dv_dec(ws, rng(C[k]), k)
    dv_list(ws, rng(C['付款账户']), '=账户列表')
    dv_fixed(ws, rng(C['审批结果']), APV_RESULTS, prompt='老板选：同意 / 部分同意（填批准金额）/ 不同意 / 暂缓；空着＝待审批')
    dv_len(ws, rng(C['单号']), 30)
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(APV_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, APV_LAST, '审批末行', N_APV)


# ───────────────────────── 采购计划 ─────────────────────────
def build_pp(ws, ctx):
    C, S = PP_COLS, PP_SHOW
    heads = [(C['日期'], '提交日期★'), (C['提交人'], '提交人'), (C['物料'], '物料及规格★'), (C['数量'], '数量'), (C['计量单位'], '单位'),
             (C['预计金额'], '预计金额'), (C['建议供应商'], '建议供应商'), (C['需用日期'], '需用日期'),
             (C['审批'], '审批\n（老板选）'), (C['采购情况'], '采购情况\n（买了选）'), (C['备注'], '备注'),
             (S['状态'], '状态\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 每天要买的东西提交一行（采购员微信报给出纳录也行）。老板在「审批」选同意 / 不同意 / 暂缓；买了在「采购情况」选已下单、已到货。'
           '这里只是计划，不算钱：到了货在【采购登记】登一行，付款走【付款审批】。【采购计划汇总】看每天的单和汇总。灰底是示例行。')
    W = {C['日期']: 11, C['提交人']: 8, C['物料']: 24, C['数量']: 9, C['计量单位']: 6, C['预计金额']: 12, C['建议供应商']: 13,
         C['需用日期']: 11, C['审批']: 9, C['采购情况']: 9, C['备注']: 20, S['状态']: 10, S['校验']: 32}
    fm = {C['日期']: DATE, C['需用日期']: DATE, C['数量']: '#,##0.###', C['预计金额']: MONEY}
    al = {C['物料']: AL, C['建议供应商']: AL, C['备注']: AL, C['数量']: AR, C['预计金额']: AR}
    rows = ctx.get('plans', [])
    _frame(ws, PP_LAST, heads, C_IN, tip, W, N_PP, fm, al, f'D{R0}', auto=tuple(S.values()), title_col='K')
    ws[f'{C["审批"]}{HDR}'].fill = fill(H_BOSS)
    for r in range(R0, R0 + N_PP):
        ws[f'{C["审批"]}{r}'].fill = fill('FFF3EAF9')
    _kpi(ws, 3, [('A', '待审批', '=COUNTIFS(购_状态码,0)&" 条"', '@'),
                 ('D', '同意了还没买', '=COUNTIFS(购_状态码,1)&" 条"', '@')])
    _nofx(ws, 'A', S['校验'], f'{C["需用日期"]}3', R0 + N_PP + 2000, f'{S["校验"]}3')
    _write_rows(ws, rows, C, date_keys=('日期', '需用日期'))
    _show(ws, N_PP, {
        S['状态']: 'INDEX(购_状态,{n})&""',
        S['校验']: _check('购'),
    }, {}, {S['校验']: AL})
    lastshow = R0 + N_PP - 1
    _check_cf(ws, S['校验'], lastshow)
    st = S['状态']
    ws.conditional_formatting.add(f'{st}{R0}:{st}{lastshow}', FormulaRule(formula=[f'{st}{R0}="待审批"'], font=F_WARN, fill=fill('FFFFF2CC')))
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_PP - 1}'
    for k in ('日期', '需用日期'):
        dv_day(ws, rng(C[k]))
    dv_list(ws, rng(C['提交人']), '=人员列表', stop=False)
    dv_dec(ws, rng(C['数量']), '数量')
    dv_dec(ws, rng(C['预计金额']), '预计金额')
    dv_list(ws, rng(C['建议供应商']), '=供应商列表', stop=False)
    dv_fixed(ws, rng(C['审批']), PP_APPROVE, prompt='老板选；空着＝待审批')
    dv_fixed(ws, rng(C['采购情况']), PP_STATES, prompt='买了选已下单、已到货；不买了选取消')
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(PP_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, PP_LAST, '计划末行', N_PP)


# ───────────────────────── 基础资料 ─────────────────────────
def build_base(ws, ctx):
    widths(ws, {'A': 18, 'B': 12, 'C': 14, 'D': 30, 'E': 3, 'F': 3, 'G': 64})
    title(ws, '基 础 资 料', 'G', C_MASTER,
          '💡 一次填好、偶尔改：① 公司名称、建账日期等；② 账户（7 个银行/微信账户＋承兑汇票，填建账日前一天晚上的余额）；'
          '③ 收支类别（资金台帐的下拉，「算到哪」决定这笔钱进利润表哪一项、冲不冲往来）；④ 人员；⑤ 每月月底估一下库存值多少（可不填）。'
          '每块上下排着，在块里插整行、删整行都只动这一块；名称不要重复。改了名字，用「查找替换」把登记表里的旧名字也改掉。')
    home_link(ws, 'G3')
    pv = ctx.get('params', {})
    put(ws, 'A3', '① 参数', F_SEC, fill(C_MASTER), align=AL)
    ws.merge_cells('A3:C3')
    notes = {'公司名称': '首页、对账单抬头用',
             '建账日': '从哪天开始用这本账：账户余额、期初应收应付都是这一天之前（前一天晚上）的数；只录这一天及以后的单子',
             '回款天数': '客户欠款多少天没回款，首页、客户往来标红提醒',
             '期初库存': '建账日仓库里料、在制品、成品（含在外协厂的料）大概值多少（按买价估）；不填就当没变',
             '月折旧': '机床等设备每月大概折旧多少，从建账月起每月算进生产成本；不填＝不算折旧'}
    for k, lbl, default in PARAMS:
        r = PAR_ROW[k]
        put(ws, f'A{r}', lbl, F_KPI_L, F_LBL, align=AL)
        ws.merge_cells(f'A{r}:B{r}')
        v = pv.get(k)
        if v in (None, ''):
            v = default
        if k == '建账日':
            v = _d(v)
        put(ws, f'C{r}', v, F_IN, FILL_IN, DATE if k == '建账日' else (MONEY if k in ('期初库存', '月折旧') else None), AC)
        put(ws, f'D{r}', notes[k], F_NOTE, align=ALW, border=False)
        ws.merge_cells(f'D{r}:G{r}')
        ws.row_dimensions[r].height = 26
    dv_day(ws, f'C{PAR_ROW["建账日"]}')
    dv_whole(ws, f'C{PAR_ROW["回款天数"]}', 1, 999, '天数')
    for k in ('期初库存', '月折旧'):
        dv_dec(ws, f'C{PAR_ROW[k]}', k)
    bn = {'账户': '收钱、付钱的地方：农行公户、工行个人、微信……期初余额填建账日前一天晚上的余额（请按网银核对）。'
                 '类型选「票据」的（承兑汇票）不算可用资金：收到承兑记进来（收客户货款），背书付货款记出去（付供应商货款），'
                 '贴现或到期托收按票面记两行内部转账转到银行，贴息另记一行「承兑贴息」。',
          '类别': '「算到哪」决定这类钱怎么算：收客户货款、付供应商货款（冲往来，要选往来单位）；生产成本、销售费用、管理费用、财务费用、税金（进利润表）；'
                 '其他收入、其他支出；内部转账（自己账户之间倒钱）、不算收支（借款还款、老板存取、保证金）、买设备（不算当月费用）。'
                 '登过采购的货付款一定选「付供应商货款」，不要再选费用，不然成本算两遍。',
          '人员': '业务员、申请人、采购计划提交人的下拉。',
          '库存': '每月月底大概估一下仓库里料、在制品、成品（含已发货没定价的、在外协厂的料）值多少，按买价估。'
                 '利润表的成本＝这个月买进来的＋生产工资电费等＋（上次估的－这次估的）。没估的月份当作没变，利润只供参考，看几个月合起来的更准。'
                 '月份填 2026-10 这样就行。'}
    heads = {'账户': (ACC_COLS, [('名称', '账户名称'), ('类型', '类型'), ('期初余额', '期初余额\n（建账日前一天）'), ('备注', '备注')],
                      ctx.get('accounts', [])),
             '类别': (CAT_COLS, [('名称', '收支类别'), ('归类', '算到哪'), ('说明', '说明')], ctx.get('cats', [])),
             '人员': (PER_COLS, [('姓名', '姓名'), ('部门', '部门'), ('备注', '备注')], ctx.get('persons', [])),
             '库存': (STK_COLS, [('月份', '月份'), ('估值', '月底库存估值'), ('备注', '备注')], ctx.get('stock', []))}
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        cols, hd, data = heads[key]
        c2 = cols[hd[-1][0]]
        put(ws, f'A{tr}', ttl, F_SEC, fill(C_MASTER), align=AL)
        ws.merge_cells(f'A{tr}:{c2}{tr}')
        put(ws, f'G{tr}', bn[key], F_NOTE, align=ALW, border=False)
        ws.merge_cells(f'G{tr}:G{hr + 4}')
        header(ws, hr, [(cols[k], t) for k, t in hd], C_MASTER, height=34)
        fm = {}
        if key == '账户':
            fm[cols['期初余额']] = MONEY
        if key == '库存':
            fm[cols['月份']] = 'yyyy-mm;@'
            fm[cols['估值']] = MONEY
        r0 = hr + 1
        style_rows(ws, r0, r0 + cap - 1, [cols[k] for k, _ in hd], fmts=fm,
                   aligns={cols[k]: AL for k in ('名称', '备注', '姓名', '说明') if k in cols})
        for i, x in enumerate(data):
            for k, _t in hd:
                v = x.get(k)
                if v not in (None, ''):
                    ws[f'{cols[k]}{r0 + i}'] = _d(v) if k == '月份' else v
        rngB = lambda k: f'{cols[k]}{r0}:{cols[k]}{r0 + cap - 1}'
        if key == '账户':
            dv_dec(ws, rngB('期初余额'))
            dv_fixed(ws, rngB('类型'), ACC_TYPES, prompt='票据（承兑汇票）不算可用资金')
        if key == '类别':
            dv_fixed(ws, rngB('归类'), KINDS, prompt='决定这类钱算到哪（见右边说明）')
        if key == '库存':
            dv_dec(ws, rngB('估值'), '估值')
        if key != '库存':
            dv_len(ws, f'A{r0}:A{r0 + cap - 1}', 30)
        put(ws, f'{c2}{r0 + cap}', f'（这一块最多 {cap} 个，超过了【数据校验】会提醒）', F_NOTE, align=AL, border=False)
    ws.freeze_panes = 'A4'
    print_setup(ws, None, landscape=False)


# ───────────────────────── 往来单位 ─────────────────────────
def build_unit(ws, ctx):
    C, S = UNIT_COLS, UNIT_SHOW
    heads = [(C['简称'], '简称★'), (C['类型'], '类型★'), (C['全称'], '全称'), (C['别名'], '其他叫法\n（几个用、隔开）'),
             (C['联系人'], '联系人'), (C['电话'], '电话'), (C['地址'], '地址'), (C['期初应收'], '期初应收\n（客户欠我们）'),
             (C['期初应付'], '期初应付\n（我们欠供应商）'), (C['期初未开票'], '期初未开票\n（可不填）'), (C['期初欠票'], '期初欠票\n（可不填）'),
             (C['开票'], '开票\n（空＝要）'), (C['默认采购类别'], '默认采购类别'), (C['备注'], '备注'),
             (S['现在'], '现在\n（自动，到截止日）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 客户、供应商都在这里建档，一家一行。简称是下拉和报表里用的名字；全称、其他叫法（快工单里的名字、个人名字）登记上，'
           '录入时写哪个都认。期初应收、期初应付填建账日前一天晚上的余额（预收、预付填负数）。不开发票的（个人供应商等）开票选「不要」。'
           '零星现买、当场付清、不用对账的不用建档。灰底两家是示例单位，正式用时连同各表的示例行一起删掉。')
    W = {C['简称']: 14, C['类型']: 11, C['全称']: 28, C['别名']: 18, C['联系人']: 8, C['电话']: 14, C['地址']: 26,
         C['期初应收']: 13, C['期初应付']: 13, C['期初未开票']: 11, C['期初欠票']: 11, C['开票']: 7, C['默认采购类别']: 11,
         C['备注']: 26, S['现在']: 22, S['校验']: 34}
    fm = {C[k]: MONEY for k in ('期初应收', '期初应付', '期初未开票', '期初欠票')}
    fm[C['电话']] = '@'
    al = {C[k]: AL for k in ('简称', '全称', '别名', '地址', '备注', '电话')}
    for k in ('期初应收', '期初应付', '期初未开票', '期初欠票'):
        al[C[k]] = AR
    rows = ctx.get('units', [])
    _frame(ws, UNIT_LAST, heads, C_MASTER, tip, W, N_UNIT, fm, al, f'B{R0}', auto=tuple(S.values()), title_col='N')
    _kpi(ws, 3, [('A', '期初应收合计', f'=SUMIFS(位_期初应收,位_是客户,1)', MONEY),
                 ('D', '期初应付合计', f'=SUMIFS(位_期初应付,位_是供应商,1)', MONEY),
                 ('G', '客户 / 供应商', '=SUM(位_是客户)&" 家 / "&SUM(位_是供应商)&" 家"', '@')])
    _write_rows(ws, rows, C, date_keys=(), text_keys=('电话',))
    yr, yp = 'INDEX(位_应收,{n})', 'INDEX(位_应付,{n})'
    isc, iss = 'INDEX(位_是客户,{n})=1', 'INDEX(位_是供应商,{n})=1'
    zr, zp = f'ABS({yr})<0.005', f'ABS({yp})<0.005'
    now = (f'IF(INDEX(位_有效,{{n}})<>1,"",IF(OR(NOT({isc}),{zr}),"",IF({yr}<0,"预收 ","应收 ")&TEXT(ABS({yr}),"#,##0.00"))'
           f'&IF(AND({isc},{iss},NOT({zr}),NOT({zp})),"；","")&IF(OR(NOT({iss}),{zp}),"",IF({yp}<0,"预付 ","应付 ")&TEXT(ABS({yp}),"#,##0.00")))')
    _show(ws, N_UNIT, {
        S['现在']: now,
        S['校验']: 'IF(INDEX(位_末行,{n})=0,"",IF(INDEX(位_校验,{n})="","✓",INDEX(位_校验,{n})))',
    }, {}, {S['现在']: AL, S['校验']: AL})
    _check_cf(ws, S['校验'], R0 + N_UNIT - 1)
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_UNIT - 1}'
    dv_fixed(ws, rng(C['类型']), UNIT_TYPES, prompt='两边都有往来的选「客户和供应商」')
    dv_fixed(ws, rng(C['开票']), UNIT_INV, prompt='空着＝要；个人供应商、不开票的客户选「不要」（不算未开票、欠票）')
    dv_fixed(ws, rng(C['默认采购类别']), PUR_TYPES, prompt='采购登记里不填采购类别时用这个')
    for k in ('期初应收', '期初应付', '期初未开票', '期初欠票'):
        dv_dec(ws, rng(C[k]), k)
    dv_len(ws, rng(C['简称']), 20)
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(UNIT_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, UNIT_LAST, '往来末行', N_UNIT)


# ───────────────────────── 固定支出 ─────────────────────────
def build_fix(ws, ctx):
    C, S = FIX_COLS, FIX_SHOW
    heads = [(C['编号'], '编号'), (C['项目'], '项目'), (C['类别'], '收支类别★'), (C['收款单位'], '收款单位'), (C['每月金额'], '每月金额★'),
             (C['每月几号'], '每月几号付'), (C['付款账户'], '付款账户'), (C['开始月份'], '开始月份\n（空＝一直）'),
             (C['结束月份'], '结束月份\n（空＝一直）'), (C['备注'], '备注'),
             (S['本月已付'], '本月已付\n（自动）'), (S['校验'], '这一行的问题\n（自动）')]
    tip = ('💡 每个月固定要付的（厂房租金、工资、社保、贷款利息……）一条一行。资金台帐付这笔钱时，关联单号选编号；'
           '不选也行：收支类别（填了收款单位的再加往来单位）对得上就自动算已付。【资金计划】看每个月还有哪些没付。'
           '季付、年付的不放这里，到时候在【付款审批】里申请。灰底是示例。')
    W = {C['编号']: 8, C['项目']: 14, C['类别']: 14, C['收款单位']: 14, C['每月金额']: 12, C['每月几号']: 9, C['付款账户']: 10,
         C['开始月份']: 11, C['结束月份']: 11, C['备注']: 22, S['本月已付']: 12, S['校验']: 40}
    fm = {C['每月金额']: MONEY, C['开始月份']: 'yyyy-mm;@', C['结束月份']: 'yyyy-mm;@', C['编号']: '@', S['本月已付']: MONEY}
    al = {C['项目']: AL, C['收款单位']: AL, C['备注']: AL, C['每月金额']: AR}
    rows = ctx.get('fixed', [])
    _frame(ws, FIX_LAST, heads, C_MASTER, tip, W, N_FIX, fm, al, f'C{R0}', auto=tuple(S.values()), title_col='J')
    _kpi(ws, 3, [('A', '本月固定支出', '=SUMIFS(固_每月金额,固_本月生效,1)', MONEY),
                 ('D', '本月已付', '=SUMIFS(固_本月已付,固_本月生效,1)', MONEY)])
    put(ws, 'G3', '="本月＝截止日所在月 "&TEXT(P_截止,"yyyy-mm")', F_NOTE, align=AL, border=False)
    _write_rows(ws, rows, C, date_keys=('开始月份', '结束月份'), text_keys=('编号',))
    _show(ws, N_FIX, {
        S['本月已付']: 'IF(INDEX(固_本月生效,{n})<>1,"",INDEX(固_本月已付,{n}))',
        S['校验']: 'IF(INDEX(固_末行,{n})=0,"",IF(INDEX(固_校验,{n})="","✓",INDEX(固_校验,{n})))',
    }, {S['本月已付']: MONEY}, {S['校验']: AL, S['本月已付']: AR})
    _check_cf(ws, S['校验'], R0 + N_FIX - 1)
    rng = lambda col: f'{col}{R0}:{col}{R0 + N_FIX - 1}'
    dv_list(ws, rng(C['类别']), '=收支类别列表')
    dv_list(ws, rng(C['收款单位']), '=往来单位列表', stop=False, prompt='可空；填了的话资金台帐的往来单位也要一样才算已付')
    dv_dec(ws, rng(C['每月金额']), '每月金额')
    dv_whole(ws, rng(C['每月几号']), 1, 31, '每月几号')
    dv_list(ws, rng(C['付款账户']), '=账户列表')
    dv_len(ws, rng(C['编号']), 12)
    _example_grey(ws, rows, R0, C['备注'])
    home_link(ws, f'{CL(CI(FIX_LAST) + 1)}1')
    print_setup(ws, f'{HDR}:{HDR}')
    _print_to_last(ws, FIX_LAST, '固定末行', N_FIX)


BUILDERS = {SH_CASH: build_cash, SH_SALE: build_sale, SH_PUR: build_pur, SH_INV: build_inv, SH_APV: build_apv, SH_PP: build_pp,
            SH_BASE: build_base, SH_UNIT: build_unit, SH_FIX: build_fix}


def build(wb, ctx):
    for name, fn in BUILDERS.items():
        fn(wb[name], ctx)
