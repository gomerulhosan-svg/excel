# -*- coding: utf-8 -*-
"""【快工单核对】（A069 补充）：月底把快工单导出的采购入库单 / 销售出库单连表头整块粘到 J5，左边自动列出每张单、
   在【采购登记】「入库单号」/【销售登记】「出库单号」里找到几行，标出「✗ 没登」的。成品入库单粘进来只提示「财务不用登」。
   粘贴区按第 5 行的表头名称找列（快工单导出列的顺序变了也认得）；续行（没有单号的行）自动跳过，所以粘贴区排了序也不怕。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Protection
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.protection import SheetProtection
from layout import *
from common import *

S = KGD_SHOW
P0, P1 = KGD_PASTE
H = 'BA'           # 隐藏帮手格：BA1 单的种类，BB1～BF1 各列在粘贴区里是第几列
HELP = dict(种类='BA', 单号列='BB', 日期列='BC', 数量列='BD', 品名列='BE', 客户列='BF')
PR = 'BG'          # 隐藏帮手列：每行有单号＝行号（BG1＝最大那个，打印区域只打到这一行）


def build(wb, ctx):
    ws = wb[SH_KGD]
    last = KGD_R0 + N_KGD - 1
    widths(ws, {S['单号']: 17, S['日期']: 11, S['品名']: 18, S['客户']: 16, S['快工单数量']: 11, S['登了几行']: 8,
                S['模板数量']: 11, S['状态']: 36, 'I': 10})
    for i in range(CI(P0), CI(P1) + 1):
        ws.column_dimensions[CL(i)].width = 13
    title(ws, '快 工 单 核 对', 'H', C_IN,
          '💡 月底用一次：快工单导出这个月的采购入库单（或销售出库单），连表头整块粘到右边 J5（右键 → 选择性粘贴 → 数值）。'
          '左边每张单自动显示在【采购登记】「入库单号」（或【销售登记】「出库单号」）里登了几行，「✗ 没登」的照磅单、送货单补登。'
          '采购入库单没有供应商、没有金额；销售出库单有客户、多数有销售总额，但没有规格，数量只是整单合计。'
          '所以都不能直接当账用，只用来查漏登。成品入库单财务不用登。'
          '左边 A～H 是公式，这张表保护着（没设密码）；右边 J～AI 可以随便清空、粘贴。')
    home_link(ws, 'I1')                                      # 放在粘贴区（J～AI）外面：清空粘贴区不会把它删掉
    hl = HELP
    hdr = f'${P0}${KGD_HDR}:${P1}${KGD_HDR}'
    has = lambda t: f'COUNTIF({hdr},"{t}")>0'
    ws[f'{hl["种类"]}1'] = (f'=IF({has("采购入库单号")},1,IF({has("销售出库单号")},2,'
                          f'IF(AND({has("入库单号")},{has("计划数量")}),3,0)))')
    kind = f'${hl["种类"]}$1'
    pos = lambda name: f'IFERROR(MATCH({name},{hdr},0),0)'
    q = lambda t: '"' + t + '"'
    ws[f'{hl["单号列"]}1'] = '=IF(' + kind + '=0,0,' + pos('CHOOSE(' + kind + ',' + q('采购入库单号') + ',' + q('销售出库单号') + ',' + q('入库单号') + ')') + ')'
    ws[f'{hl["日期列"]}1'] = '=' + pos('IF(' + kind + '=2,' + q('出库日期') + ',' + q('入库日期') + ')')
    ws[f'{hl["数量列"]}1'] = '=' + pos('IF(' + kind + '=2,' + q('出库数量') + ',' + q('入库总数') + ')')
    ws[f'{hl["品名列"]}1'] = f'={pos(chr(34) + "商品名称" + chr(34))}'
    ws[f'{hl["客户列"]}1'] = f'={pos(chr(34) + "客户" + chr(34))}'
    for k, c in hl.items():
        ws[f'{c}2'] = k
        ws[f'{c}1'].font = ws[f'{c}2'].font = F_HELP
        hide(ws, c)
    col = lambda k: f'${hl[k]}$1'

    # 第 3 行：认出来是哪种单、几张、没登几张；第 4 行：怎么用
    a_rng, h_rng = f'${S["单号"]}${KGD_R0}:${S["单号"]}${last}', f'${S["状态"]}${KGD_R0}:${S["状态"]}${last}'
    n_pre = f'COUNTIF({h_rng},"⚠ 建账前*")'
    summary = (f'=IF({kind}=0,IF(COUNTA({hdr})=0,"把快工单导出的表（连表头）粘到 J5：右键 → 选择性粘贴 → 数值",'
               f'"⚠ 第 5 行认不出是哪种单（要连表头粘到 J5；快工单格式变了？）"),'
               f'CHOOSE({kind},"采购入库单","销售出库单","成品入库单")&"：共 "&COUNTIF({a_rng},"?*")&" 张"&IF({kind}=3,"（成品入库财务不用登）",'
               f'"，没登 "&COUNTIF({h_rng},"✗*")&" 张"'
               f'&IF({kind}=2,"，数量对不上 "&COUNTIF({h_rng},"⚠ 已登*")&" 张","")'          # 采购入库单不比数量（吨/公斤），不写这一项
               f'&IF({n_pre}>0,"，⚠ 建账前的单却登了 "&{n_pre}&" 张","")'
               f'&"，建账前 "&COUNTIF({h_rng},"建账前*")&" 张"))'
               # 左边筛过（第 ④ 步）没放开：清空、粘贴只动看得见的行，会清不干净、粘错位
               f'&IF(SUBTOTAL(103,{a_rng})<ROWS({a_rng}),"　⚠ 左边有筛选：先全部显示再清空、粘贴","")')
    put(ws, 'A3', summary, Font(name=YH, sz=12, bold=True, color='FF1F3864'), fill('FFFFF2CC'), align=ALW)
    ws.merge_cells('A3:H3')
    ws.row_dimensions[3].height = 36
    ws.conditional_formatting.add('A3', FormulaRule(formula=['OR(ISNUMBER(SEARCH("⚠",A3)),AND(ISNUMBER(SEARCH("没登",A3)),ISERROR(SEARCH("没登 0 张",A3))))'],
                                                    font=Font(name=YH, sz=12, bold=True, color='FFC00000')))
    put(ws, 'A4', ('步骤：① 快工单导出这个月的单子，连表头整块复制；'
                   f'② 先 数据→筛选→全部显示（上次左边筛过的要先放开），再选中 {P0}{KGD_HDR}:{P1}{last} 按 Delete，清掉上次的；'
                   '③ 点 J5 → 右键「选择性粘贴」→「数值」；④ 左边「状态」筛「✗ 没登」，照磅单、送货单补登。'
                   '钢材快工单是公斤、模板是吨，数量差 1000 倍是正常的。一次只粘一种单。左边用筛选，别排序（排了也不会变）。'),
        F_NOTE, align=ALW, border=False)
    ws.merge_cells('A4:H4')
    ws.row_dimensions[4].height = 50
    header(ws, KGD_HDR, [(S['单号'], '单号'), (S['日期'], '日期'), (S['品名'], '商品名称（第一样）'), (S['客户'], '客户'),
                         (S['快工单数量'], '快工单数量\n（整单合计）'), (S['登了几行'], '登了几行'),
                         (S['模板数量'], '登的数量合计'), (S['状态'], '状态')], C_IN, height=34)
    put(ws, f'I{KGD_HDR}', None, border=False)
    for i in range(CI(P0), CI(P1) + 1):                      # 粘贴区表头：淡蓝
        c = ws.cell(row=KGD_HDR, column=i)
        c.fill = FILL_PASTE
        c.font = F_TXTB
        c.alignment = ACW
    put(ws, f'I{KGD_HDR - 1}', '粘贴区 → 从 J5 开始粘（连表头）', F_RED, align=AL, border=False)    # 写在粘贴区外面，清空 J～AI 不会删掉

    # 每一行
    for r in range(KGD_R0, last + 1):
        row = f'${P0}{r}:${P1}{r}'
        cell = lambda k: f'INDEX({row},1,{col(k)})'
        A = f'${S["单号"]}{r}'
        ws[f'{S["单号"]}{r}'] = f'=IF(OR({kind}=0,{col("单号列")}=0),"",TRIM({cell("单号列")}&""))'
        ws[f'{S["日期"]}{r}'] = f'=IF(OR({A}="",{col("日期列")}=0),"",IFERROR({date_parse(cell("日期列"))[0:]},""))'
        ws[f'{S["品名"]}{r}'] = f'=IF(OR({A}="",{col("品名列")}=0),"",TRIM({cell("品名列")}&""))'
        ws[f'{S["客户"]}{r}'] = f'=IF(OR({A}="",{kind}<>2,{col("客户列")}=0),"",TRIM({cell("客户列")}&""))'
        ws[f'{S["快工单数量"]}{r}'] = f'=IF(OR({A}="",{col("数量列")}=0),"",IFERROR(--{cell("数量列")},""))'
        key = f'"*"&{esc(A)}&"*"'
        ws[f'{S["登了几行"]}{r}'] = (f'=IF({A}="","",IF({kind}=1,COUNTIF(采_采购单号,{key}),IF({kind}=2,COUNTIF(销_出库单号,{key}),"")))')
        ws[f'{S["模板数量"]}{r}'] = (f'=IF({A}="","",IF({kind}=1,SUMIFS(采_数量,采_采购单号,{key}),'
                                  f'IF({kind}=2,SUMIFS(销_数量,销_出库单号,{key}),"")))')
        B, E, F, G = (f'${S[k]}{r}' for k in ('日期', '快工单数量', '登了几行', '模板数量'))
        ws[f'{S["状态"]}{r}'] = (f'=IF({A}="","",IF({kind}=3,"成品入库：财务不用登",IF(NOT(ISNUMBER({B})),"⚠ 日期看不懂",'
                               f'IF(AND({B}<P_建账日,N({F})>0),"⚠ 建账前的单却登了 "&{F}&" 行：期初欠款里可能已经有了",'
                               f'IF({B}<P_建账日,"建账前（不用登）",IF({F}=0,"✗ 没登",'
                               f'IF(AND({kind}=2,ISNUMBER({E}),ABS(N({E})-N({G}))>0.001),"⚠ 已登 "&{F}&" 行，数量差 "&(ROUND({E}-{G},3)&""),'
                               f'"✓ 已登 "&{F}&" 行")))))))')
        ws[f'{PR}{r}'] = f'=IF({A}="",0,ROW())'
        for k in S:
            c = ws[f'{S[k]}{r}']
            c.font, c.fill, c.border = F_AUTO, FILL_AUTO, BD
            c.alignment = AL if k in ('单号', '品名', '客户', '状态') else (AC if k in ('日期', '登了几行') else AR)
        ws[f'{S["日期"]}{r}'].number_format = DATE
        ws[f'{S["快工单数量"]}{r}'].number_format = 'General'
        ws[f'{S["模板数量"]}{r}'].number_format = 'General'
    rs = f'{S["状态"]}{KGD_R0}:{S["状态"]}{last}'
    ws.conditional_formatting.add(rs, FormulaRule(formula=[f'LEFT({S["状态"]}{KGD_R0},1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(rs, FormulaRule(formula=[f'LEFT({S["状态"]}{KGD_R0},1)="⚠"'],
                                                  font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.conditional_formatting.add(rs, FormulaRule(formula=[f'LEFT({S["状态"]}{KGD_R0},1)="✓"'], font=Font(name=YH, sz=10, color='FF70AD47')))
    # 预置的快工单导出（真实数据，当普通值写）
    for i, rowv in enumerate(ctx.get('kgd', []) or []):
        for j, v in enumerate(rowv):
            if v in (None, ''):
                continue
            c = ws.cell(row=KGD_HDR + i, column=CI(P0) + j)
            if isinstance(v, str) and v[:1] in '=+-@':
                c.value = "'" + v
            else:
                c.value = v
            if i:
                c.font = F_TXT
    ws.freeze_panes = f'{P0}{KGD_R0}'
    ws.auto_filter.ref = f'A{KGD_HDR}:{S["状态"]}{last}'
    print_setup(ws, f'{KGD_HDR}:{KGD_HDR}')
    # 打印只到最后一张单（BG1＝最后一行有单号的行号；空着＝只打表头），不打一千行空框
    ws[f'{PR}1'] = f'=MAX({KGD_HDR},MAX({PR}{KGD_R0}:{PR}{last}))'
    ws[f'{PR}2'] = '打印到第几行'
    ws[f'{PR}1'].font = ws[f'{PR}2'].font = F_HELP
    hide(ws, PR)
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${S["状态"]}$1:${S["状态"]}${last},{q}!${PR}$1)')
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'

    # 保护：左边 A～H 的公式、帮手格不让改，也不能删列、插删行（删列会让整张表 #REF!）；
    # 粘贴区 J～AI 整列不锁（列样式＋已经写过的格子），清空、粘贴照样能用；筛选、改格式、调列宽行高都允许
    unlock = Protection(locked=False)
    for i in range(CI(P0), CI(P1) + 1):
        ws.column_dimensions[CL(i)].protection = unlock
    for c in list(ws._cells.values()):
        if CI(P0) <= c.column <= CI(P1):
            c.protection = unlock
    ws.protection = SheetProtection(sheet=True, autoFilter=False, sort=False, formatColumns=False, formatRows=False,
                                    formatCells=False, selectLockedCells=False, selectUnlockedCells=False)
