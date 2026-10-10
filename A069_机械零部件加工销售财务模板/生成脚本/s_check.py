# -*- coding: utf-8 -*-
"""【数据校验】：各登记表逐行问题（✗ 要改、⚠ 要看）、基础资料重名、对数检查和该跟进的事，最后列出有问题的行（表名＋行号＋问题）。
   首页「提醒」引用 CHK_X / CHK_W / CHECK_ROWS。"""
from openpyxl.formatting.rule import FormulaRule
from layout import *
from common import *

CHK_X, CHK_W = 'C4', 'F4'
CHECK_ROWS = {}
D, Y = 'P_截止', 'P_截止年月'

# 逐行检查的表：(录入表, 隐藏表前缀, 有内容的标志列（录入表）, 显示「问题」的列, 容量, 常见原因)
SHEETS = [
    (SH_CASH, '资', CASH_COLS['日期'], CASH_SHOW['校验'], N_CASH,
     '没填日期或金额、日期写成 10.8、账户或类别不在基础资料里、收付货款没选往来单位、关联单号找不到'),
    (SH_SALE, '销', SALE_COLS['日期'], SALE_SHOW['校验'], N_SALE, '客户不在往来单位表、日期早于建账日、还没定价（⚠）'),
    (SH_PUR, '采', PUR_COLS['日期'], PUR_SHOW['校验'], N_PUR, '供应商不在往来单位表、日期早于建账日、还没定价（⚠）'),
    (SH_INV, '票', INV_COLS['日期'], INV_SHOW['校验'], N_INV, '单位不在往来单位表、两头往来的单位没选开出/收到、发票号码重复'),
    (SH_APV, '批', APV_COLS['申请日期'], APV_SHOW['校验'], N_APV, '没填单号或单号重复、部分同意没填批准金额、老板没同意却付了（⚠）'),
    (SH_PP, '购', PP_COLS['日期'], PP_SHOW['校验'], N_PP, '没填物料、审批或采购情况乱填、到了需用日期还没买（⚠）'),
    (SH_UNIT, '位', UNIT_COLS['简称'], UNIT_SHOW['校验'], N_UNIT, '简称重复、没选类型、全称跟别家简称一样'),
    (SH_FIX, '固', FIX_COLS['类别'], FIX_SHOW['校验'], N_FIX, '类别不在基础资料里、编号重复、类别和收款单位跟上面一条一样（⚠）'),
]


def _over(sh, cap, col='A'):
    """录入表第 cap 条以后还有没有东西（区域从表头行锚定，插删行跟着动）"""
    return f'COUNTA(INDEX({sh}!${col}:${col},ROW({sh}!${col}${HDR})+{cap + 1}):INDEX({sh}!${col}:${col},1048576))'


def _over_base():
    """基础资料每块：第 cap 个以后、下一块标题以前还有没有名字（块变短时 X位<=cap，不算超）"""
    parts = []
    for key, ttl, tr, hr, cap in BASE_BLOCKS:
        a = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+{cap + 1})'
        b = f'INDEX({SH_BASE}!$A:$A,ROW({SH_BASE}!$A${hr})+P_{key}位)'
        parts.append(f'IF(P_{key}位<={cap},0,COUNTA({a}:{b}))')
    return '=' + '+'.join(parts)


def _nofx(sh, key_col, chk_col, cap):
    last = R0 + cap + 2000
    return f'COUNTIFS({sh}!${key_col}${R0}:${key_col}${last},"<>",{sh}!${chk_col}${R0}:${chk_col}${last},"")'


def build(wb, ctx=None):
    ws = wb[SH_CHK]
    widths(ws, {'A': 6, 'B': 52, 'C': 14, 'D': 7, 'E': 12, 'F': 14, 'G': 58})
    title(ws, '数 据 校 验', 'G', C_CHK,
          '💡 录完数看一眼这张：✗＝一定要改（不改报表会少算或算错），⚠＝看一下对不对、该跟进了，ℹ＝说明一下（不算问题）。'
          '先看每张表有几条问题，再看对数检查，最后按「表＋行号」回到那张表找到改掉（每张登记表最右边「这一行的问题」也写着）。这张表全是自动的。')
    home_link(ws, 'G3')
    put(ws, 'B4', '要改的（✗）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    put(ws, 'E4', '要看的（⚠）合计', F_KPI_L, fill('FFD9E1F2'), align=AC)
    xs, wsum = [], []

    # ① 各登记表
    r = 6
    section(ws, r, 'A', 'G', '① 各登记表逐行检查（每行最右边「这一行的问题」汇总）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '表'), ('C', '✗ 要改'), ('D', ''), ('E', '⚠ 要看'), ('F', '右边没公式'), ('G', '常见原因')], C_CHK)
    for i, (sh, p, kc, cc, cap, why) in enumerate(SHEETS):
        rr = r + 2 + i
        nm = f'{p}_校验'
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        link(put(ws, f'B{rr}', sh, Font(name=YH, sz=10, bold=True, color='FF0563C1', underline='single'), align=AL), sh)
        put(ws, f'C{rr}', f'=COUNTIF({nm},"✗*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'D{rr}', f'=IF(C{rr}>0,"✗","✓")', F_TXTB, align=AC)
        put(ws, f'E{rr}', f'=COUNTIF({nm},"⚠*")', F_AUTOB, fmt='0', align=AC)
        put(ws, f'F{rr}', f'={_nofx(sh, kc, cc, cap)}', F_AUTOB, fmt='0;-0;"✓"', align=AC)
        put(ws, f'G{rr}', why, F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 30
        xs.append(f'C{rr}')
        wsum.append(f'E{rr}')
        wsum.append(f'IF(F{rr}>0,1,0)')
        CHECK_ROWS[sh] = rr
    put(ws, f'B{r + 2 + len(SHEETS)}', '「右边没公式」：插进来的行最右边几格是空的（不影响汇总，只是那一行看不到余额和问题）：选上一行右边灰色几格往下拉一下就有了。',
        F_NOTE, align=ALW, border=False)
    ws.merge_cells(f'B{r + 2 + len(SHEETS)}:G{r + 2 + len(SHEETS)}')
    ws.row_dimensions[r + 2 + len(SHEETS)].height = 28
    r += len(SHEETS) + 4

    # ② 基础资料
    section(ws, r, 'A', 'G', '② 基础资料（名字不能重复、类别要选「算到哪」）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '个数'), ('D', '状态'), ('G', '怎么改')], C_CHK)
    hcols = {'账户': ('AA', '账户_名称', N_ACC), '类别': ('AB', '类别_名称', N_CAT), '人员': ('AC', '人员_姓名', N_PER)}
    for key, (hc, rng, n) in hcols.items():
        for i in range(n):
            nm = f'INDEX({rng},{i + 1})'
            ws[f'{hc}{i + 1}'] = f'=IF({nm}="",0,IF(COUNTIF({rng},{esc(nm)})>1,1,0))'
            ws[f'{hc}{i + 1}'].font = F_HELP
    kinds_ok = 'OR(' + ','.join(f'{{kd}}="{k}"' for k in KINDS) + ')'
    for i in range(N_CAT):                                   # 类别有名字没归类（或归类写错）
        nm, kd = f'INDEX(类别_名称,{i + 1})', f'INDEX(类别_归类,{i + 1})'
        ws[f'AD{i + 1}'] = f'=IF({nm}="",0,IF({kinds_ok.format(kd=kd)},0,1))'
        ws[f'AD{i + 1}'].font = F_HELP
    for i in range(N_STK):                                   # 库存估值：月份重复
        k = f'INDEX(库存_年月,{i + 1})'
        ws[f'AE{i + 1}'] = f'=IF({k}=0,0,IF(COUNTIF(库存_年月,{k})>1,1,0))'
        ws[f'AE{i + 1}'].font = F_HELP
    for i in range(N_STK):                                   # 库存估值：月份填了但看不懂
        raw = f'INDEX({inref(SH_BASE, STK_COLS["月份"], block="库存")},{i + 2})'
        ws[f'AG{i + 1}'] = f'=IF(AND(INDEX(库存_年月,{i + 1})=0,TRIM({raw}&"")<>""),1,0)'
        ws[f'AG{i + 1}'].font = F_HELP
    hide(ws, 'AA', 'AB', 'AC', 'AD', 'AE', 'AG')
    checks = [('账户名称重复', f'=SUM($AA$1:$AA${N_ACC})', '✗', '同一个账户只写一行'),
              ('收支类别名称重复', f'=SUM($AB$1:$AB${N_CAT})', '✗', '同一个类别只写一行'),
              ('人员姓名重复', f'=SUM($AC$1:$AC${N_PER})', '✗', '重名的加个字区分'),
              ('收支类别没选「算到哪」（或者乱填）', f'=SUM($AD$1:$AD${N_CAT})', '✗', '在【基础资料】③ 收支类别的「算到哪」列从下拉选'),
              ('一个账户都没有', '=IF(COUNTIF(账户_名称,"?*")=0,1,0)', '✗', '在【基础资料】② 账户 至少填一个'),
              ('没有「算到哪」是收客户货款 / 付供应商货款的类别', '=IF(COUNTIF(类别_归类,"收客户货款")=0,1,0)+IF(COUNTIF(类别_归类,"付供应商货款")=0,1,0)',
               '✗', '回款、付货款要有对应的类别，往来才冲得掉'),
              ('往来单位简称重复', '=COUNTIF(位_校验,"✗ 简称重复*")', '✗', '同一家只建一行；别的叫法写在「其他叫法」里'),
              ('库存估值同一个月填了两次', f'=SUM($AE$1:$AE${N_STK})', '✗', '一个月只填一个估值'),
              ('库存估值的月份看不懂', f'=SUM($AG$1:$AG${N_STK})', '⚠', '月份填 2026-10 这样（Excel 会显示成 2026-10）'),
              ]
    for i, (lab, f, kind, how) in enumerate(checks):
        rr = r + 2 + i
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        put(ws, f'B{rr}', lab, F_TXT, align=AL)
        put(ws, f'C{rr}', f, F_AUTOB, fmt='0', align=AC)
        put(ws, f'D{rr}', f'=IF(C{rr}>0,"{kind}","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        (xs if kind == '✗' else wsum).append(f'IF(C{rr}>0,1,0)')
        CHECK_ROWS[lab] = rr
    r += len(checks) + 3

    # ③ 对数检查、该跟进的
    section(ws, r, 'A', 'G', '③ 对数检查、该跟进的（截至首页截止日期）', C_CHK)
    header(ws, r + 1, [('A', '#'), ('B', '检查'), ('C', '金额/个数'), ('D', '状态'), ('G', '说明')], C_CHK)
    for i in range(N_ACC):                                   # 每个账户截至截止日的余额 < 0 ？
        a = f'INDEX(账户_名称,{i + 1})'
        ws[f'AF{i + 1}'] = (f'=IF({a}="",0,IF(INDEX(账户_期初余额,{i + 1})+SUMIFS(资_净额,资_账户号,{i + 1},资_日期,"<="&{D})<-0.005,1,0))')
        ws[f'AF{i + 1}'].font = F_HELP
    hide(ws, 'AF')
    upto = lambda nm: f'{nm},"<="&{D}'
    overdue = '位_是客户,1,位_没回款天数,">"&P_回款天数'
    ex = '+'.join(f'COUNTIF({sh}!${col}:${col},"示例*")' for sh, col in (
        (SH_CASH, CASH_COLS['备注']), (SH_SALE, SALE_COLS['备注']), (SH_PUR, PUR_COLS['备注']), (SH_INV, INV_COLS['备注']),
        (SH_APV, APV_COLS['备注']), (SH_PP, PP_COLS['备注']), (SH_UNIT, UNIT_COLS['备注']), (SH_FIX, FIX_COLS['备注'])))
    late_fix_n = (f'SUMPRODUCT((固_本月生效=1)*(固_每月几号<DAY({D}))*(固_每月几号>0)*(固_本月已付<固_每月金额-0.005))')
    num = [
        ('内部转账没配对（转出、转进合计不为 0）', f'=ROUND(SUMIFS(资_净额,资_归类,"内部转账",{upto("资_日期")}),2)', 'ABS(C{r})>=0.01', '✗',
         '自己账户之间倒钱要记两行（转出账户一行支出、转进账户一行收入），金额一样；资金台帐按类别筛「内部转账」查'),
        ('账户余额是负数的账户个数（按日期算到截止日）', f'=SUM($AF$1:$AF${N_ACC})', 'C{r}>0', '⚠', '钱不够付还付出去了？多半漏记了收款或内部转进，或者期初余额没填'),
        ('日期比上面早的行（即时余额按上下顺序算）', '=COUNTIF(资_日期乱,1)', 'C{r}>0', '⚠', '资金台帐整表按日期从小到大排一下序，右边的即时余额就对了（报表都按日期算，不受影响）'),
        ('有日期打错的嫌疑（最后几笔比别的晚一个多月）',
         f'=IF(AND(NOT(ISNUMBER({SH_HOME}!{HOME_END})),P_截止被挡=1),P_近月笔数,0)', 'C{r}>0', '⚠',
         '截止日期没跟着这几笔走（还按之前最后一笔算）；资金台帐、销售、采购按日期从大到小排一下就找到，多半年份填错了'),
        ('首页截止日期比最后一笔早', f'=IF(AND(ISNUMBER({SH_HOME}!{HOME_END}),{D}<P_最后日期),P_最后日期-{D},0)', 'C{r}>0', '⚠',
         '首页黄格填了截止日期，以后录的就不算进来；看完了记得清空'),
        ('日期晚于截止日期的行（暂不算）', f'=COUNTIFS(资_有效,1,资_日期,">"&{D})+COUNTIFS(销_有效,1,销_日期,">"&{D})+COUNTIFS(采_有效,1,采_日期,">"&{D})'
                               f'+COUNTIFS(票_有效,1,票_日期,">"&{D})', 'C{r}>0', '⚠', '还没到日子的不算；年份填错的改过来'),
        ('示例行还没删（灰底、备注写着「示例」）', f'={ex}', 'C{r}>0', '⚠', '正式用之前，把各张登记表、往来单位、固定支出里的示例行整行删掉'),
        ('="客户超过 "&P_回款天数&" 天没回款（家数）"', f'=COUNTIFS({overdue})', 'C{r}>0', '⚠', '看【客户往来】，催款；提醒天数在【基础资料】改'),
        ('还没定价的发货（笔）', f'=COUNTIFS(销_待定价,1,销_计入往来,1,销_日期,"<="&{D})', 'C{r}>0', '⚠', '销售登记里单价、金额空着的：这些货的钱还没算进应收和收入，定了价回到原行补上'),
        ('还没定价的到货（笔）', f'=COUNTIFS(采_待定价,1,采_计入往来,1,采_日期,"<="&{D})', 'C{r}>0', '⚠', '采购登记里单价、金额空着的：还没算进欠供应商和成本'),
        ('等老板审批的付款（金额）', '=SUMIFS(批_申请金额,批_状态码,0)', 'C{r}>0.005', 'ℹ', '看【待审批付款单】，打印给老板签'),
        ('老板批了还没付的（金额）', '=SUM(批_未付)', 'C{r}>0.005', 'ℹ', '看【资金计划】'),
        ('本月固定支出过了日子还没付清（条）', f'={late_fix_n}', 'C{r}>0', '⚠', '看【资金计划】①；付了以后资金台帐关联单号选编号（或者类别、收款单位对上）'),
        ('预付了钱的供应商（家）', '=COUNTIFS(位_是供应商,1,位_应付,"<-0.005")', 'C{r}>0', '⚠', '先付款后到货的正常；到了货要记采购登记'),
        ('采购计划到了需用日期还没买（条）', '=COUNTIF(购_校验,"⚠ 已经到了需用日期*")', 'C{r}>0', '⚠', '看【采购计划汇总】'),
        ('上个月月底没估库存', f'=IF(COUNTIFS(库存_有值,1)=0,0,IF(COUNTIFS(库存_年月,IF(MOD({Y},100)=1,{Y}-89,{Y}-1),库存_有值,1)=0,1,0))',
         'C{r}>0', 'ℹ', '不估也行：那个月的利润只供参考，看几个月合计更准'),
        ('基础资料没填建账日期', f'=IF(ISNUMBER({SH_BASE}!$C${PAR_ROW["建账日"]}),0,1)', 'C{r}>0', '⚠', '期初余额、期初欠款都是建账日前一天晚上的数，一定要填'),
        ('基础资料没填公司名称', '=IF(P_公司名称="",1,0)', 'C{r}>0', 'ℹ', '首页、对账单抬头用'),
        (f'资金台帐超过 {N_CASH} 笔（多出来的不算）', '=' + _over(SH_CASH, N_CASH), 'C{r}>0', '✗', '一本最多这么多：另起一本（账户期初余额、往来期初填上一本的余额）'),
        (f'销售登记超过 {N_SALE} 行', '=' + _over(SH_SALE, N_SALE), 'C{r}>0', '✗', '同上'),
        (f'采购登记超过 {N_PUR} 行', '=' + _over(SH_PUR, N_PUR), 'C{r}>0', '✗', '同上'),
        (f'发票登记超过 {N_INV} 张', '=' + _over(SH_INV, N_INV), 'C{r}>0', '✗', '同上'),
        (f'付款审批超过 {N_APV} 单', '=' + _over(SH_APV, N_APV, APV_COLS['申请日期']), 'C{r}>0', '✗', '同上'),
        (f'采购计划超过 {N_PP} 行', '=' + _over(SH_PP, N_PP), 'C{r}>0', '✗', '同上'),
        (f'往来单位超过 {N_UNIT} 家', '=' + _over(SH_UNIT, N_UNIT), 'C{r}>0', '✗', '不常来往的合并或删掉'),
        (f'固定支出超过 {N_FIX} 条', '=' + _over(SH_FIX, N_FIX, FIX_COLS['类别']), 'C{r}>0', '✗', '合并几条'),
        ('基础资料某一块超了（多出来的不算）', _over_base(), 'C{r}>0', '✗', f'账户最多 {N_ACC} 个、收支类别 {N_CAT} 个、人员 {N_PER} 个、库存估值 {N_STK} 个月'),
        (f'待付款的审批单超过 {LINK_N} 张（下拉列不全）', f'=MAX(0,COUNTIF(批_下拉标签,"?*")+COUNTIFS(固_有效,1,固_编号,"?*")-{LINK_N})', 'C{r}>0', 'ℹ',
         '关联单号也可以直接打单号'),
    ]
    r += 2
    for i, (lab, f, bad, kind, how) in enumerate(num):
        rr = r + i
        put(ws, f'B{rr}', lab, F_TXT, align=ALW)
        put(ws, f'A{rr}', i + 1, F_TXT, align=AC)
        isn = any(w in lab for w in ('个数', '（笔）', '（条）', '（家', '行', '超', '没填', '比最后', '晚于', '嫌疑', '月底', '家数', '张', '单'))
        put(ws, f'C{rr}', f, F_AUTOB, fmt='0' if isn and '金额' not in lab else MONEY, align=AR)
        put(ws, f'D{rr}', f'=IF({bad.format(r=rr)},"{kind}","✓")', F_TXTB, align=AC)
        put(ws, f'G{rr}', how, F_NOTE, align=ALW)
        ws.row_dimensions[rr].height = 30
        if kind == '✗':
            xs.append(f'IF({bad.format(r=rr)},1,0)')
        elif kind == '⚠':
            wsum.append(f'IF({bad.format(r=rr)},1,0)')
        CHECK_ROWS[lab] = rr
    r += len(num) + 1
    put(ws, CHK_X, '=' + '+'.join(xs), F_KPI_V, fill('FFFFFFFF'), '0', AC)
    put(ws, CHK_W, '=' + '+'.join(wsum), F_KPI_V, fill('FFFFFFFF'), '0', AC)

    # ④ 有问题的行
    section(ws, r, 'A', 'G', '④ 有问题的行（回到那张表按行号找；只列 ✗ 和 ⚠）', C_CHK)
    header(ws, r + 1, [('A', '第几条'), ('B', '问题'), ('C', '表'), ('D', '行号'), ('E', '日期'), ('F', '金额'), ('G', '内容')], C_CHK)
    r += 2
    lists = [
        (SH_CASH, '资', N_CASH, 100, '资_日期', 'INDEX(资_收入,{k})-INDEX(资_支出,{k})',
         'INDEX(资_账户,{k})&"｜"&INDEX(资_类别,{k})&"｜"&INDEX(资_单位,{k})&"｜"&INDEX(资_显示摘要,{k})'),
        (SH_SALE, '销', N_SALE, 60, '销_日期', 'INDEX(销_金额,{k})', 'INDEX(销_单位,{k})&"｜"&INDEX(销_品名规格,{k})'),
        (SH_PUR, '采', N_PUR, 60, '采_日期', 'INDEX(采_金额,{k})', 'INDEX(采_单位,{k})&"｜"&INDEX(采_品名规格,{k})'),
        (SH_INV, '票', N_INV, 40, '票_日期', 'INDEX(票_价税合计,{k})', 'INDEX(票_方向,{k})&"｜"&INDEX(票_单位,{k})&"｜"&INDEX(票_发票号码,{k})'),
        (SH_APV, '批', N_APV, 40, '批_申请日期', 'INDEX(批_申请金额,{k})', 'INDEX(批_单号,{k})&"｜"&INDEX(批_单位,{k})&"｜"&INDEX(批_付款内容,{k})'),
        (SH_PP, '购', N_PP, 30, '购_日期', 'INDEX(购_预计金额,{k})', 'INDEX(购_物料,{k})'),
        (SH_UNIT, '位', N_UNIT, 30, None, '0', 'INDEX(位_简称,{k})'),
        (SH_FIX, '固', N_FIX, 20, None, 'INDEX(固_每月金额,{k})', 'INDEX(固_编号,{k})&"｜"&INDEX(固_项目,{k})'),
    ]
    cc_i = CI('AH')
    for sh, p, n, show, date_nm, amt, desc in lists:
        chk, row_nm = f'{p}_校验', f'{p}_录入行'
        cc = CL(cc_i)
        cc_i += 1
        counter(ws, cc, 1, n, lambda i, chk=chk: f'OR(LEFT(INDEX({chk},{i + 1}),1)="✗",LEFT(INDEX({chk},{i + 1}),1)="⚠")')
        hide(ws, cc)
        total = f'${cc}${n}'
        put(ws, f'B{r}', f'="{sh}：共 "&{total}&" 行有问题"&IF({total}>{show},"（只列前 {show} 行）","")', F_TXTB, FILL_SUB, align=AL)
        ws.merge_cells(f'B{r}:G{r}')
        r += 1
        for k in range(1, show + 1):
            rr = r + k - 1
            ix = f'$A{rr}'
            ws[f'A{rr}'] = f'=IF({kth(k, cc, 1, n)}=0,"",{kth(k, cc, 1, n)})'
            ws[f'A{rr}'].font = F_HELP
            g = lambda e: f'=IF({ix}="","",{e.format(k=ix)})'
            put(ws, f'B{rr}', g(f'INDEX({chk},{{k}})'), F_TXT, align=AL)
            put(ws, f'C{rr}', f'=IF({ix}="","","{sh}")', F_TXT, align=AC)
            put(ws, f'D{rr}', g(f'INDEX({row_nm},{{k}})'), F_TXTB, fmt='0', align=AC)
            if date_nm:
                put(ws, f'E{rr}', g(f'IF(INDEX({date_nm},{{k}})=0,"",INDEX({date_nm},{{k}}))'), F_TXT, fmt=DATE, align=AC)
            else:
                put(ws, f'E{rr}', None, F_TXT, align=AC)
            put(ws, f'F{rr}', g(amt), F_TXT, fmt=MONEY, align=AR)
            put(ws, f'G{rr}', g(desc), F_NOTE, align=AL)
        r += show + 1
    for col in ('D',):
        ws.conditional_formatting.add(f'{col}1:{col}{r}', FormulaRule(formula=[f'{col}1="✗"'], font=F_RED))
        ws.conditional_formatting.add(f'{col}1:{col}{r}', FormulaRule(formula=[f'{col}1="⚠"'],
                                                                     font=Font(name=YH, sz=10, bold=True, color='FFC65911')))
    ws.conditional_formatting.add(f'B1:B{r}', FormulaRule(formula=['LEFT(B1,1)="✗"'], font=F_RED))
    ws.conditional_formatting.add(f'B1:B{r}', FormulaRule(formula=['LEFT(B1,1)="⚠"'], font=Font(name=YH, sz=10, color='FFC65911')))
    ws.freeze_panes = 'A5'
    print_setup(ws, '1:4', landscape=False)
    ws.print_area = f'A1:G{r}'
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
