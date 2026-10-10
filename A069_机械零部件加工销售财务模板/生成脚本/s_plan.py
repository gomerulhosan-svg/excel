# -*- coding: utf-8 -*-
"""计划组（紫）：资金计划、待审批付款单、采购计划汇总。只用 layout 的定义名称和本表格子，不引用录入表。口径见 agent_common「口径速查」。

【资金计划】黄格 B3 月份 m（2026-10 / 202610 / 只填 10＝截止日那年的 10 月；空＝P_截止年月）。
   第 5～12 行（固定位置，首页可引用）：C5 现在可用资金（账户_可用=1，按日期算到 P_截止）、C6 本月固定支出还没付（① 没付合计；
   m 比截止年月晚时再加截止月还没付的、截止月以后到 m 前一个月整月的，标签变成「含之前月份」）、
   C7 批了还没付（② 还差合计）、C8 预计还剩＝C5－C6－C7、C9 等老板审批（③ 合计）、C10 如果待审批的都批了＝C8－C9、
   C11 手上票据（承兑，账户_可用=0，不算可用）、C12 采购计划同意了还没买（购_状态码=1 的 购_预计金额，只供参考）。
   第 14 行起四块一块接一块往下排（前一块有几行占几行，不留大段空行）：
   ① 固定支出（固_有效=1、开始年月≤m≤结束年月，按 _固 顺序）：已付＝按编号键 ＋（这个月先占的）按 "固|"&匹配键，资_年月=m、资_日期≤P_截止；
      「这个月先占」＝m 月生效、本表上面 m 月生效的行里没有同一匹配键（不用 固_先占：换了新的一条时旧的不生效，付款要算给新的）；
      没付＝MAX(0,应付－已付)；状态：已付清 / 没付 / 付了一部分；m＝截止年月、0<每月几号<截止日的日、没付清 →「过了日子还没付」。
   ② 批了还没付（批_状态码 1/2、批_计划年月≤m，按 批_排序键，最多 100 单）；③ 等老板审批（状态码 0、计划年月≤m，最多 100 单）；
   ④ 暂缓（状态码 5，全部，最多 50 单）。合计都是全部的（不受「只列前几条」影响）。②③ 提示：没填/过了计划日期、
   收款单位跟本月生效的固定支出一样（可能重复登记）、③ 这家有到货没定价。「回首页」在 K3（打印范围 A～J 外）。
【待审批付款单】打印给老板签字（A4 横向一页宽，第 7 行表头每页重复）。第 5 行（固定位置）：A5 现在可用资金、D5 手上票据、
   F5 本月（截止月）固定支出还没付、H5 已批未付（全部）、J5 待审批合计、L5 都批了还剩（＝A5－F5－H5－J5）。
   第 8 行起往下排：待审批（状态码 0，按 批_排序键，最多 100 单）→ 合计 → 暂缓（状态码 5，最多 50 单）→ 合计 → 老板签字。
   超过显示条数时，清单下面的合计只算列出来的（老板签字的单子上看得到的），全部的在第 5 行。
   「批示」「批准金额」两栏不放数，只画框给老板手写。
【采购计划汇总】黄格 C3 日期 d（空＝P_截止）、F3 月份 m（空＝P_截止年月）、H3 打印（当天的单 / 全部；空＝当天的单，只印 ①）。
   第 7 行起往下排：① d 那天提交的（购_排序键＝提交日期×SORT_M＋n 连着的一段，最多 100 条）；
   ② m 月提交的按状态（7 种）、按提交人（人员_姓名 顺序，同名只算第一个，外加「其他」＝合计－各人）；
   ③ 待审批、④ 同意没买（全部，按需用日期，没填的排最后，最多 100 条；需用日期≤P_截止 标红「已到日子」）。
   C3 空着、采购计划最晚的提交日期比截止日晚时，A4 提醒去选那天。冻结到第 4 行；「回首页」在 K3（打印范围外）。
隐藏列（三张表一样）：AA 名字、AB 值（标量，第 3 行起）；AC～AG 往下排的「段、段内第几行、行类、第几条、来源行」；AI～AP 每段的分界（第 3 行起）；
   AR～BA、BF 固定支出逐条 / AR～AV 人员逐个（第 3 行起）；BB～BE 账户（第 3 行起 15 行）；BH 起排序键辅助列（第 2 行起，每条源数据一行）；
   BL～BR 资金计划的截止月固定支出（选以后的月份用）；counter() 的条件列从 CA 往右。
行类：1 段标题、2 表头、5 明细、3 合计、6 共几条、7x 附加行（签字）、0 空。"""
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from layout import *
from common import *

M = SORT_M
PURPLE_H = 'FF8064A2'                 # 表头（浅一点的紫）
_thin = Side(style='thin', color='FFBFBFBF')
CF_BD = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
SIGN_BD = Border(bottom=Side(style='thin', color='FF000000'))
LBL = fill('FFE4DFEC')
F_KV = Font(name=YH, sz=11, bold=True, color='FF1F3864')
F_KVB = Font(name=YH, sz=12, bold=True, color='FF1F3864')
YM_FMT = '[<13]0"月";[<100000]yyyy-mm;0'   # 只填 10 显示「10月」；日期显示 2026-10；202610 原样
NUM = 'General'
SHOW_APV, SHOW_HOLD, SHOW_PP = 100, 50, 100
ST_PP = ['待审批', '同意没买', '已下单', '已到货', '不同意', '暂缓', '取消']
ST_PP_NOTE = ['老板还没批', '批了还没下单', '', '', '', '先放一放', '']


def cf_fill(rgb):
    return PatternFill('solid', fgColor=rgb, bgColor=rgb)


def R2(x):
    return f'ROUND({x},2)'


class Sc:
    """隐藏标量：AA 列名字、AB 列值（第 3 行起）"""

    def __init__(self, ws, keys, r0=3, cl='AA', cv='AB'):
        self.ws, self.cv = ws, cv
        self.r = {k: r0 + i for i, k in enumerate(keys)}
        for k, r in self.r.items():
            ws[f'{cl}{r}'] = k
            ws[f'{cl}{r}'].font = F_HELP

    def __call__(self, k):
        return f'${self.cv}${self.r[k]}'

    def set(self, k, f):
        c = self.ws[f'{self.cv}{self.r[k]}']
        c.value = '=' + f
        c.font = F_HELP


def hset(ws, coord, f):
    ws[coord] = f if isinstance(f, (int, float)) else '=' + f
    ws[coord].font = F_HELP


def note(ws, coord, last, f, height=None, font=F_NOTE):
    r = ''.join(ch for ch in coord if ch.isdigit())
    ws.merge_cells(f'{coord}:{last}{r}')
    put(ws, coord, f, font, align=ALW, border=False)
    if height:
        ws.row_dimensions[int(r)].height = height


def cf_blocks(ws, r0, r1, rules):
    """条件格式：rules＝[(起列, 止列, 公式(本块第一列字母), 样式)]，按优先顺序；
       列切成「规则集合一样」的连续块，每格只落在一个块里（LibreOffice / 有的 WPS 对重叠的块只认一个）。"""
    lo = min(CI(a) for a, _b, _c, _k in rules)
    hi = max(CI(b) for _a, b, _c, _k in rules)
    runs = []
    for i in range(lo, hi + 1):
        sig = tuple(j for j, (a, b, _c, _k) in enumerate(rules) if CI(a) <= i <= CI(b))
        if not sig:
            continue
        if runs and runs[-1][2] == sig and runs[-1][1] == i - 1:
            runs[-1][1] = i
        else:
            runs.append([i, i, sig])
    for i0, i1, sig in runs:
        for j in sig:
            _a, _b, cond, kw = rules[j]
            ws.conditional_formatting.add(f'{CL(i0)}{r0}:{CL(i1)}{r1}',
                                          FormulaRule(formula=[cond(CL(i0))], stopIfTrue=True, **kw))


def month_any(x):
    """真日期 → 年月；202610 → 本身；看不懂 → 0"""
    return (f'IF(AND({x}>=200001,{x}<=209912,INT({x})={x},MOD({x},100)>=1,MOD({x},100)<=12),{x},'
            f'IF(AND({x}>=36526,{x}<73051),YEAR({x})*100+MONTH({x}),0))')


def month_selector(ws, sc, cell, keys=('输入', '认出前', 'm', '认出', 'Y', 'mo')):
    """年月黄格：2026-10（变成日期）/ 202610 / 只填 1～12（截止日那年）；空或看不懂＝P_截止年月"""
    k_in, k_pre, k_m, k_ok, k_y, k_mo = keys
    v = sc(k_in)
    sc.set(k_in, f'IF(ISNUMBER({cell}),{cell},IFERROR(--TRIM({cell}&""),0))')
    sc.set(k_pre, f'IF(AND({v}>=1,{v}<=12,INT({v})={v}),YEAR(P_截止)*100+{v},{month_any(v)})')
    sc.set(k_m, f'IF(TRIM({cell}&"")="",P_截止年月,IF({sc(k_pre)}=0,P_截止年月,{sc(k_pre)}))')
    sc.set(k_ok, f'IF(TRIM({cell}&"")="",1,IF({sc(k_pre)}>0,1,0))')
    sc.set(k_y, f'INT({sc(k_m)}/100)')
    sc.set(k_mo, f'MOD({sc(k_m)},100)')
    dv = DataValidation(type='custom', formula1='TRUE', allow_blank=True, showErrorMessage=False, showInputMessage=True,
                        promptTitle='提示', prompt='填 2026-10 或 202610，只填 10 也行（截止日那年）；空着＝截止日那个月')
    ws.add_data_validation(dv)
    dv.add(cell.replace('$', ''))


def neg(v, word):
    """负数写成文字：预付 / 多收票"""
    return f'IF({v}="","",IF({v}<-0.005,"{word} "&TEXT(-{v},"#,##0.00"),{v}))'


def dt0(v):
    return f'IF({v}>0,{v},"")'


# ───────────────────────── 往下排的几段清单 ─────────────────────────
class Flow:
    """从第 r0 行起几段清单一块接一块往下排。每段：[标题] [表头] 明细 [合计] [共几条] [附加行…] [空行]。
       隐藏列 AC～AG：段、段内第几行、行类、第几条、来源行；AI～AP：每段的 起、标题止、表头止、明细止、合计止、共几条止、附加止、段长。
       段的定义：title/head/tot/trail＝{列: 公式}（不带等号）；item＝{列: fn(来源行, 第几条)}；extra＝[{列: 公式}…]；
       n_all＝条数（格子）；cap＝最多列几条；src＝fn(第几条)→来源行（>0）。"""

    def __init__(self, ws, sc, r0, cols, hcols=('AC', 'AD', 'AE', 'AF', 'AG'), th='AI', th_r0=3):
        self.ws, self.sc, self.r0, self.cols = ws, sc, r0, cols
        self.H = dict(zip(('段', '段内', '类', '第几', '源'), hcols))
        self.thc = [CL(CI(th) + i) for i in range(8)]
        self.th_r0 = th_r0
        self.segs = []

    def add(self, **kw):
        for k in ('title', 'head', 'tot', 'trail'):
            kw.setdefault(k, None)
        kw.setdefault('extra', [])
        kw.setdefault('spacer', 1)
        kw.setdefault('item', {})
        self.segs.append(kw)

    @staticmethod
    def size(s):
        return ((1 if s['title'] else 0) + (1 if s['head'] else 0) + s['cap'] + (1 if s['tot'] else 0)
                + (1 if s['trail'] else 0) + len(s['extra']) + s['spacer'])

    def th(self, ci, i):
        return f'${self.thc[ci]}${self.th_r0 + i}'

    def rng(self, ci):
        return f'${self.thc[ci]}${self.th_r0}:${self.thc[ci]}${self.th_r0 + len(self.segs) - 1}'

    def row_of(self, i, ci):
        """第 i 段（0 起）分界 ci 那一行的前一行（比如 ci=5：共几条那一行）"""
        return f'{self.r0}+{self.th(0, i)}+{self.th(ci, i)}-1'

    def build(self, fmts, aligns, font=F_TXT, height=None, shrink=(), wrap=()):
        ws, S = self.ws, len(self.segs)
        self.n = n = sum(self.size(s) for s in self.segs)
        for ci, t in enumerate(('起', '标题止', '表头止', '明细止', '合计止', '共几条止', '附加止', '段长')):
            ws[f'{self.thc[ci]}{self.th_r0 - 1}'] = t
            ws[f'{self.thc[ci]}{self.th_r0 - 1}'].font = F_HELP
        for i, s in enumerate(self.segs):
            r = self.th_r0 + i
            na, th = s['n_all'], self.th
            hset(ws, f'{self.thc[0]}{r}', '0' if i == 0 else f'{th(0, i - 1)}+{th(7, i - 1)}')
            hset(ws, f'{self.thc[1]}{r}', str(1 if s['title'] else 0))
            hset(ws, f'{self.thc[2]}{r}', f'{th(1, i)}+' + (f'IF({na}>0,1,0)' if s['head'] else '0'))
            hset(ws, f'{self.thc[3]}{r}', f'{th(2, i)}+MAX(0,MIN({na},{s["cap"]}))')
            hset(ws, f'{self.thc[4]}{r}', f'{th(3, i)}+' + (f'IF({na}>0,1,0)' if s['tot'] else '0'))
            hset(ws, f'{self.thc[5]}{r}', f'{th(4, i)}+{1 if s["trail"] else 0}')
            hset(ws, f'{self.thc[6]}{r}', f'{th(5, i)}+{len(s["extra"])}')
            hset(ws, f'{self.thc[7]}{r}', f'{th(6, i)}+{s["spacer"]}')
        self.end = f'({self.th(0, S - 1)}+{self.th(7, S - 1)})'
        for k, c in self.H.items():
            ws[f'{c}{self.r0 - 1}'] = k
            ws[f'{c}{self.r0 - 1}'].font = F_HELP
        nx = max(len(s['extra']) for s in self.segs)
        for j in range(n):
            r = self.r0 + j
            sg, o, t, k, src = (f'${self.H[x]}{r}' for x in ('段', '段内', '类', '第几', '源'))
            f = f'IF({j}<{self.end},{S},0)'
            for i in range(S - 1, 0, -1):
                f = f'IF({j}<{self.th(0, i)},{i},{f})'
            hset(ws, sg[1:], f)
            ix = lambda ci: f'INDEX({self.rng(ci)},{sg})'
            hset(ws, o[1:], f'IF({sg}=0,0,{j}-{ix(0)})')
            hset(ws, t[1:], (f'IF({sg}=0,0,IF({o}<{ix(1)},1,IF({o}<{ix(2)},2,IF({o}<{ix(3)},5,IF({o}<{ix(4)},3,'
                             f'IF({o}<{ix(5)},6,IF({o}<{ix(6)},70+{o}-{ix(5)},0)))))))'))
            hset(ws, k[1:], f'IF({t}=5,{o}-{ix(2)}+1,0)')
            hset(ws, src[1:], f'IF({k}=0,0,{self.by_seg([s["src"](k) for s in self.segs], sg)})')
            for c in self.cols:
                parts = [(5, self.by_seg([s['item'][c](src, k) if c in s['item'] else None for s in self.segs], sg))]
                for code, key in ((1, 'title'), (2, 'head'), (3, 'tot'), (6, 'trail')):
                    parts.append((code, self.by_seg([(s[key] or {}).get(c) for s in self.segs], sg)))
                for e in range(nx):
                    parts.append((70 + e, self.by_seg([s['extra'][e].get(c) if e < len(s['extra']) else None
                                                       for s in self.segs], sg)))
                out = '""'
                for code, e in reversed([(x, e) for x, e in parts if e != '""']):
                    out = f'IF({t}={code},{e},{out})'
                cell = ws[f'{c}{r}']
                cell.value = '=' + out
                cell.font = font
                al = aligns.get(c, AC)
                cell.alignment = Alignment(horizontal=al.horizontal, vertical='center', shrink_to_fit=(c in shrink),
                                           wrap_text=(c in wrap))       # 长文字自动换行（缩小字体太长时看不清）
                if c in fmts:
                    cell.number_format = fmts[c]
            if height:
                ws.row_dimensions[r].height = height
        self.r1 = self.r0 + n - 1

    @staticmethod
    def by_seg(exprs, sg):
        vals = [(i + 1, e) for i, e in enumerate(exprs) if e not in (None, '""')]
        if not vals:
            return '""'
        if len(vals) == len(exprs) and all(e == vals[0][1] for _i, e in vals):
            return vals[0][1]
        out = '""'
        for i, e in reversed(vals):
            out = f'IF({sg}={i},{e},{out})'
        return out

    def cf(self, last, extra_rules=(), title_color=C_PLAN, head_color=PURPLE_H):
        T = lambda: f'${self.H["类"]}{self.r0}'
        rules = list(extra_rules) + [
            ('A', last, lambda c: f'{T()}=1', dict(fill=cf_fill(title_color), font=Font(bold=True, color='FFFFFFFF'))),
            ('A', last, lambda c: f'{T()}=2', dict(fill=cf_fill(head_color), font=Font(bold=True, color='FFFFFFFF'), border=CF_BD)),
            ('A', last, lambda c: f'{T()}=3', dict(fill=cf_fill('FFFCE4D6'), font=Font(bold=True), border=CF_BD)),
            ('A', last, lambda c: f'{T()}=6', dict(font=Font(color='FF808080'))),
            ('A', last, lambda c: f'{T()}=5', dict(border=CF_BD)),
        ]
        cf_blocks(self.ws, self.r0, self.r1, rules)

    def t(self):
        return f'${self.H["类"]}{self.r0}'

    def s(self):
        return f'${self.H["段"]}{self.r0}'


# ───────────────────────── 隐藏小表：账户、固定支出、排序辅助 ─────────────────────────
def acct_table(ws, r0=3, cols=('BB', 'BC', 'BD', 'BE')):
    """账户第 a 个（第 r0+a-1 行）：名称、第一次出现、可用、截止日余额（按日期）；返回（可用资金、手上票据）公式"""
    cn, cf_, ca, cb = cols
    for c, t in zip(cols, ('账户', '第一次', '可用', '截止余额')):
        ws[f'{c}{r0 - 1}'] = t
        ws[f'{c}{r0 - 1}'].font = F_HELP
    for i in range(N_ACC):
        a, r = i + 1, r0 + i
        hset(ws, f'{cn}{r}', f'INDEX(账户_名称,{a})&""')
        hset(ws, f'{cf_}{r}', f'IF({cn}{r}="",0,IF(IFERROR(MATCH({esc(f"{cn}{r}")},账户_名称,0),0)={a},1,0))')
        hset(ws, f'{ca}{r}', f'IF({cf_}{r}=1,INDEX(账户_可用,{a}),0)')
        hset(ws, f'{cb}{r}', f'IF({cf_}{r}<>1,0,ROUND(INDEX(账户_期初余额,{a})+SUMIFS(资_净额,资_账户号,{a},资_日期,"<="&P_截止),2))')
    g = lambda c: f'${c}${r0}:${c}${r0 + N_ACC - 1}'
    return (R2(f'SUMIFS({g(cb)},{g(cf_)},1,{g(ca)},1)'), R2(f'SUMIFS({g(cb)},{g(cf_)},1,{g(ca)},0)'))


FIX_C = dict(生效='AR', 应付='AS', 已付='AT', 没付='AU', 过期='AV', 状态='AW', 计数='AX', 键='AY', 先占='AZ', 单位='BA',
             名='BF')
FIX_CUT = dict(生效='BL', 应付='BM', 已付='BN', 没付='BO', 键='BP', 先占='BQ', 后月='BR')   # 资金计划：截止月（选以后的月份时用）


def mon_idx(x):
    """YYYYMM → 月序号（年×12＋月），好算两个年月之间隔几个月"""
    return f'(INT({x}/100)*12+MOD({x},100))'


def fix_table(ws, m, r0=3, C=FIX_C, gate=None, lite=False):
    """固定支出第 i 条（第 r0+i-1 行）在月份 m：生效、应付、已付（资_年月=m、≤截止）、没付、过期、状态；返回汇总公式 dict。
       先占按「这个月」算：这个月生效、而且本表上面这个月生效的行里没有同一个匹配键（换了新的一条、旧的不生效了，
       没填关联单号的付款就算给新的那条）。gate：另加的生效条件；lite：只算到没付（不要状态、计数）。"""
    for k, c in C.items():
        ws[f'{c}{r0 - 1}'] = k
        ws[f'{c}{r0 - 1}'].font = F_HELP
    upto = '资_日期,"<="&P_截止'
    last = r0 + N_FIX - 1
    for i in range(N_FIX):
        n, r = i + 1, r0 + i
        g = lambda k: f'INDEX(固_{k},{n})'
        on, due, paid, left, key, first = (f'{C[k]}{r}' for k in ('生效', '应付', '已付', '没付', '键', '先占'))
        cond = f'{g("有效")}=1,{g("开始年月")}<={m},{g("结束年月")}>={m}' + (f',{gate}' if gate else '')
        hset(ws, on, f'IF(AND({cond}),1,0)')
        hset(ws, key, f'IF({on}=1,{g("匹配键")},"")')
        prev = (f'COUNTIFS(${C["键"]}${r0}:{C["键"]}{r - 1},{esc(key)},${C["生效"]}${r0}:{C["生效"]}{r - 1},1)'
                if i else '0')
        hset(ws, first, f'IF({on}<>1,0,IF({prev}=0,1,0))')
        hset(ws, due, f'IF({on}=1,{g("每月金额")},0)')
        k1 = f'IF({g("编号键")}="",0,SUMIFS(资_支出额,资_归属,{esc(g("编号键"))},资_年月,{m},{upto}))'
        k2 = (f'IF({first}=1,SUMIFS(资_支出额,资_归属,{esc(chr(34) + "固|" + chr(34) + "&" + key)},'
              f'资_年月,{m},{upto}),0)')
        hset(ws, paid, f'IF({on}<>1,0,ROUND({k1}+{k2},2))')
        hset(ws, left, f'IF({on}<>1,0,MAX(0,ROUND({due}-{paid},2)))')
        if lite:
            continue
        late = f'{C["过期"]}{r}'
        hset(ws, late, f'IF(AND({on}=1,{left}>0.005,{m}=P_截止年月,{g("每月几号")}>0,{g("每月几号")}<DAY(P_截止)),1,0)')
        hset(ws, f'{C["状态"]}{r}', (f'IF({on}<>1,"",IF({late}=1,"过了日子还没付",IF({left}<=0.005,"已付清",'
                                    f'IF({paid}>0.005,"付了一部分","没付"))))'))
        # 给 ②③ 提示「可能跟固定支出重复」：这个月生效、填了收款单位的，放收款单位和编号（没编号放项目）
        un = f'{C["单位"]}{r}'
        hset(ws, un, f'IF({on}=1,{g("收款单位")}&"","")')
        hset(ws, f'{C["名"]}{r}', f'IF({un}="","",IF({g("编号")}<>"",{g("编号")},{g("项目")}))')
    g = lambda k: f'${C[k]}${r0}:${C[k]}${last}'
    out = dict(没付=R2(f'SUMIFS({g("没付")},{g("生效")},1)'), rng=g)
    if lite:
        return out
    counter(ws, C['计数'], r0, N_FIX, lambda i: f'${C["生效"]}${r0 + i}=1')
    out.update(n=cnt(C['计数'], r0, N_FIX), 应付=R2(f'SUMIFS({g("应付")},{g("生效")},1)'),
               已付=R2(f'SUMIFS({g("已付")},{g("生效")},1)'),
               过期n=f'COUNTIFS({g("过期")},1)', 未清n=f'COUNTIFS({g("生效")},1,{g("没付")},">0.005")',
               kth=lambda k: kth(k, C['计数'], r0, N_FIX))
    return out


def fix_dup(fx, unit):
    """②③ 提示：收款单位跟这个月生效的某条固定支出一样 →「；可能跟固定支出 GD04 重复」，不然空"""
    return (f'IF({unit}="","",IFERROR("；可能跟固定支出 "&INDEX({fx["rng"]("名")},MATCH({esc(unit)},{fx["rng"]("单位")},0))'
            f'&" 重复",""))')


def key_col(ws, col, n_src, cond, key, title):
    """排序辅助列：第 i 条源数据（第 i+1 行）满足条件＝排序键，不然空；返回区域"""
    ws[f'{col}1'] = title
    ws[f'{col}1'].font = F_HELP
    for i in range(n_src):
        hset(ws, f'{col}{i + 2}', f'IF({cond(i + 1)},{key(i + 1)},"")')
    return f'${col}$2:${col}${n_src + 1}'


def finish(ws, last, end_cell, cap_row, freeze, hdr_rows=None, landscape=True):
    ws.freeze_panes = freeze
    print_setup(ws, hdr_rows, landscape=landscape)
    ws.oddFooter.center.text = '第 &P 页 共 &N 页'
    q = f"'{ws.title}'"
    ws.defined_names['Print_Area'] = DefinedName(
        'Print_Area', attr_text=f'{q}!$A$1:INDEX({q}!${last}$1:${last}${cap_row},{q}!{end_cell})')


def kpi_rows(ws, r, last, items):
    """r 行左边标签（A:B 合并）、C 值、D:last 灰字说明"""
    for i, (lab, f, nt, strong) in enumerate(items):
        rr = r + i
        ws.merge_cells(f'A{rr}:B{rr}')
        put(ws, f'A{rr}', lab, F_KPI_L, LBL, align=AL)
        put(ws, f'B{rr}', None, F_KPI_L, LBL)
        put(ws, f'C{rr}', '=' + f, F_KVB if strong else F_KV, FILL_TOT if strong else FILL_AUTO, MONEY, AR)
        note(ws, f'D{rr}', last, nt)
        ws.row_dimensions[rr].height = 21


# ═══════════════════════════════ 资金计划 ═══════════════════════════════
TIP_PLAN = ('💡 看这个月的钱够不够付。黄格选月份（填 2026-10 或 202610，只填 10 也行；空着＝截止日那个月）。'
            '顶上：现在可用资金（银行、微信、现金……算到截止日）减掉这个月固定支出还没付的、老板批了还没付的，就是预计还剩'
            '（选以后的月份，截止日那个月起还没付的固定支出也一起扣）；再减掉等老板审批的，看都批了够不够。'
            '下面四块：① 这个月的固定支出（【固定支出】里这个月生效的；付款时资金台帐关联单号选编号，'
            '或者类别、收款单位对上，就算付了；已付按付款月份算到截止日）；② 老板批了还没付完的；③ 等老板审批的（②③ 都是计划在这个月及以前的；'
            '「现在欠这家」＝我们还欠收款单位的货款，「预付」＝先付了钱，「多收票」＝对方发票开多了）；④ 暂缓的（全部）。')
FP_R0 = 14


def build_fplan(ws):
    LAST = 'J'
    widths(ws, {'A': 13, 'B': 16, 'C': 22, 'D': 14, 'E': 12.5, 'F': 13.5, 'G': 13.5, 'H': 13.5, 'I': 11, 'J': 22, 'K': 10})
    title(ws, '=IF(P_公司名称="","",P_公司名称&"　")&"资金计划"', LAST, C_PLAN, TIP_PLAN)
    selector(ws, 'A3', '选月份', 'B3', None, fmt=YM_FMT)
    home_link(ws, 'K3')                       # 放在打印范围（A～J）外面
    ws.row_dimensions[3].height = 30
    keys = ['输入', '认出前', 'm', '认出', 'Y', 'mo', '可用', '票据', '固n', '固应付', '固已付', '固没付', '固过期n', '固未清n',
            '批n', '批批准', '批已付', '批还差', '审n', '审申请', '缓n', '缓申请', '购n', '购金额', '预计还剩', '都批了还剩',
            '尾', '末行', '之前没付', '固没付含前']
    sc = Sc(ws, keys)
    month_selector(ws, sc, 'B3')
    m, Y, mo = sc('m'), sc('Y'), sc('mo')
    avail, bills = acct_table(ws)
    sc.set('可用', avail)
    sc.set('票据', bills)
    fx = fix_table(ws, m)
    for k in ('应付', '已付', '没付', '过期n', '未清n'):
        sc.set(f'固{k}', fx[k])
    sc.set('固n', fx['n'])
    # 选以后的月份：截止日那个月还没付的 ＋ 中间几个月（截止月以后的月份还没有付款，整月都算没付）
    later = f'{m}>P_截止年月'
    fxc = fix_table(ws, 'P_截止年月', C=FIX_CUT, gate=later, lite=True)
    C2 = FIX_CUT
    for i in range(N_FIX):
        n = i + 1
        g = lambda k: f'INDEX(固_{k},{n})'
        span = (f'MAX(0,MIN({mon_idx(g("结束年月"))},{mon_idx(m)}-1)-MAX({mon_idx(g("开始年月"))},{mon_idx("P_截止年月")}+1)+1)')
        hset(ws, f'{C2["后月"]}{3 + i}', f'IF(AND({later},{g("有效")}=1),ROUND({g("每月金额")}*{span},2),0)')
    sc.set('之前没付', f'IF({later},ROUND({fxc["没付"]}+SUM(${C2["后月"]}$3:${C2["后月"]}${N_FIX + 2}),2),0)')
    sc.set('固没付含前', R2(f'{sc("固没付")}+{sc("之前没付")}'))
    # 排序辅助：② 批了没付、③ 待审批、④ 暂缓
    pi = lambda k, n: f'INDEX(批_{k},{n})'
    K2 = key_col(ws, 'BH', N_APV, lambda n: (f'AND({pi("有效", n)}=1,OR({pi("状态码", n)}=1,{pi("状态码", n)}=2),'
                                              f'{pi("计划年月", n)}<={m})'), lambda n: pi('排序键', n), '②键')
    K3 = key_col(ws, 'BI', N_APV, lambda n: f'AND({pi("有效", n)}=1,{pi("状态码", n)}=0,{pi("计划年月", n)}<={m})',
                 lambda n: pi('排序键', n), '③键')
    K4 = key_col(ws, 'BJ', N_APV, lambda n: f'AND({pi("有效", n)}=1,{pi("状态码", n)}=5)', lambda n: pi('排序键', n), '④键')
    ok12 = lambda nm: (f'SUMIFS({nm},批_有效,1,批_状态码,1,批_计划年月,"<="&{m})'
                       f'+SUMIFS({nm},批_有效,1,批_状态码,2,批_计划年月,"<="&{m})')
    sc.set('批n', f'COUNT({K2})')
    sc.set('批批准', R2(ok12('批_批准金额')))
    sc.set('批已付', R2(ok12('批_已付')))
    sc.set('批还差', R2(ok12('批_未付')))
    sc.set('审n', f'COUNT({K3})')
    sc.set('审申请', R2(f'SUMIFS(批_申请金额,批_有效,1,批_状态码,0,批_计划年月,"<="&{m})'))
    sc.set('缓n', f'COUNT({K4})')
    sc.set('缓申请', R2('SUMIFS(批_申请金额,批_有效,1,批_状态码,5)'))
    sc.set('购n', 'COUNTIFS(购_有效,1,购_状态码,1)')
    sc.set('购金额', R2('SUMIFS(购_预计金额,购_有效,1,购_状态码,1)'))
    sc.set('预计还剩', R2(f'{sc("可用")}-{sc("固没付含前")}-{sc("批还差")}'))
    sc.set('都批了还剩', R2(f'{sc("预计还剩")}-{sc("审申请")}'))

    ws.merge_cells(f'C3:I3')
    put(ws, 'C3', (f'="实际用的："&{Y}&" 年 "&{mo}&" 月"&IF(TRIM(B3&"")="","（空着＝截止日那个月）","")'
                   f'&IF({sc("认出")}=0,"　⚠ 没认出这个月份，先按截止日那个月","")'
                   f'&IF({later},"　（这个月还没到：钱只算到截止日 "&TEXT(P_截止,"yyyy-mm-dd")&"，固定支出都算没付；'
                   f'截止日那个月起还没付的固定支出也一起扣）","")'
                   f'&IF({m}<P_截止年月,"　（以前的月份：已付算到截止日）","")'
                   f'&IF(DATE({Y},{mo}+1,0)<P_建账日,"　⚠ 这个月在建账以前，没有流水（固定支出都算没付）","")'),
        F_NOTE, align=ALW, border=False)
    ws.conditional_formatting.add('C3', FormulaRule(formula=['ISNUMBER(FIND("⚠",$C$3))'], font=Font(bold=True, color='FFC00000')))

    # ── 顶上：够不够 ──
    section(ws, 4, 'A', LAST, f'="这个月的钱够不够（"&{Y}&" 年 "&{mo}&" 月；钱算到截止日 "&TEXT(P_截止,"yyyy-mm-dd")&"）"', C_PLAN)
    left = sc('预计还剩')
    allok = sc('都批了还剩')
    cy, cmo = 'INT(P_截止年月/100)', 'MOD(P_截止年月,100)'
    prev_m = f'IF({mo}=1,({Y}-1)&" 年 12 月",{Y}&" 年 "&({mo}-1)&" 月")'
    kpi_rows(ws, 5, LAST, [
        ('现在可用资金', sc('可用'), '="银行、微信、支付宝、现金等账户截至 "&TEXT(P_截止,"yyyy-mm-dd")&" 的余额合计（不含承兑汇票）"', False),
        (f'=IF({later},"－ 固定支出还没付（含之前月份）","－ 本月固定支出还没付")', sc('固没付含前'),
         f'="下面 ① 的「没付」合计"&IF({later},"（"&TEXT({sc("固没付")},"#,##0.00")&"）","")&"：这个月 "&{sc("固n")}&" 条固定支出，还有 "'
         f'&{sc("固未清n")}&" 条没付清"&IF({sc("固过期n")}>0,"（"&{sc("固过期n")}&" 条过了日子）","")'
         f'&IF({later},"；另加 "&{cy}&" 年 "&{cmo}&" 月（截止日那个月）"'
         f'&IF({mon_idx(m)}-{mon_idx("P_截止年月")}=1,"","到 "&{prev_m})&"还没付的 "&TEXT({sc("之前没付")},"#,##0.00"),"")',
         False),
        ('－ 批了还没付', sc('批还差'), f'="下面 ② 的「还差」合计："&{sc("批n")}&" 单（老板同意了、计划在这个月及以前、还没付完的）"', False),
        ('＝ 预计还剩', left, f'=IF({left}<-0.005,"⚠ 不够付，还差 "&TEXT(-{left},"#,##0.00")&"：先催回款、贴现承兑，或者跟老板商量哪些晚点付","够付")', True),
        ('－ 等老板审批的', sc('审申请'), f'="下面 ③ 合计："&{sc("审n")}&" 单（计划在这个月及以前的申请金额）"', False),
        ('＝ 如果待审批的都批了', allok, f'=IF({allok}<-0.005,"⚠ 都批了就不够，还差 "&TEXT(-{allok},"#,##0.00"),"都批了也够")', True),
        ('另：手上票据（承兑汇票）', sc('票据'), '不算可用资金：可以背书转给供应商付货款，或者贴现到银行', False),
        ('另：采购计划同意了还没买的', sc('购金额'), f'="【采购计划】老板同意了、还没下单的预计金额（"&{sc("购n")}&" 条），只供参考，没算进上面"', False),
    ])
    cf = ws.conditional_formatting
    for rr in (8, 10):
        cf.add(f'C{rr}', FormulaRule(formula=[f'C{rr}<-0.005'], font=Font(bold=True, color='FFC00000')))
        cf.add(f'D{rr}', FormulaRule(formula=[f'LEFT(D{rr},1)="⚠"'], font=Font(bold=True, color='FFC00000')))
    cf.add('D6', FormulaRule(formula=[f'{sc("固过期n")}>0'], font=Font(bold=True, color='FFC00000')))

    # ── ①②③④ 往下排 ──
    fl = Flow(ws, sc, FP_R0, 'ABCDEFGHIJ')
    G = lambda k, n: f'INDEX(固_{k},{n})'
    FT = lambda k, n: f'INDEX({fx["rng"](k)},{n})'
    fl.add(title=dict(A='"① 固定支出"', B=f'{Y}&" 年 "&{mo}&" 月"', C='"这个月该付的"'),
           head=dict(A='"编号"', B='"项目"', C='"收支类别"', D='"收款单位"', E='"每月几号"', F='"应付"', G='"已付"',
                     H='"没付"', I='"付款账户"', J='"状态"'),
           item=dict(A=lambda n, k: G('编号', n), B=lambda n, k: G('项目', n), C=lambda n, k: G('类别', n),
                     D=lambda n, k: G('收款单位', n), E=lambda n, k: f'IF({G("每月几号", n)}>0,{G("每月几号", n)}&" 号","")',
                     F=lambda n, k: FT('应付', n), G=lambda n, k: FT('已付', n), H=lambda n, k: FT('没付', n),
                     I=lambda n, k: G('付款账户', n), J=lambda n, k: FT('状态', n)),
           tot=dict(A='"合计"', F=sc('固应付'), G=sc('固已付'), H=sc('固没付')),
           trail=dict(A=f'IF({sc("固n")}=0,"（没有）","共 "&{sc("固n")}&" 条")',
                      C=f'IF({sc("固n")}=0,"这个月没有固定支出",IF({sc("固过期n")}>0,"⚠ "&{sc("固过期n")}&" 条过了日子还没付",""))'),
           n_all=sc('固n'), cap=N_FIX, src=lambda k: fx['kth'](k))
    P = lambda k, n: f'INDEX(批_{k},{n})'
    # 提示：几样都有就用「；」连起来（每样前面带「；」，最后去掉第一个）
    plan_hint = lambda n: (f'IF({P("计划付款日期", n)}=0,"；没填计划日期",IF({P("计划付款日期", n)}<P_截止,"；过了计划日期",""))')
    dup = lambda n: fix_dup(fx, P('单位', n))
    pend = lambda n: (f'IF({P("单位号", n)}=0,"",IF(INDEX(位_待定价到货,MAX(1,{P("单位号", n)}))>0,'
                      f'"；另有 "&INDEX(位_待定价到货,MAX(1,{P("单位号", n)}))&" 笔到货没定价（没算进欠款）",""))')
    common = dict(A=lambda n, k: P('单号', n), B=lambda n, k: P('单位', n), C=lambda n, k: P('付款内容', n),
                  D=lambda n, k: dt0(P('申请日期', n)), E=lambda n, k: dt0(P('计划付款日期', n)))
    fl.add(title=dict(A='"② 批了没付"', B='"计划在本月及以前"', C='"老板同意了、还没付完的"'),
           head=dict(A='"单号"', B='"收款单位"', C='"付款内容"', D='"申请日期"', E='"计划付款日期"', F='"批准金额"', G='"已付"',
                     H='"还差"', I='"付款账户"', J='"提示"'),
           item=dict(common, F=lambda n, k: P('批准金额', n), G=lambda n, k: P('已付', n), H=lambda n, k: P('未付', n),
                     I=lambda n, k: P('付款账户', n), J=lambda n, k: f'MID({plan_hint(n)}&{dup(n)},2,200)'),
           tot=dict(A='"合计"', F=sc('批批准'), G=sc('批已付'), H=sc('批还差')),
           trail=dict(A=f'IF({sc("批n")}=0,"（没有）","共 "&{sc("批n")}&" 单")',
                      C=f'IF({sc("批n")}>{SHOW_APV},"只列了前 {SHOW_APV} 单（合计是全部的）","")'),
           n_all=sc('批n'), cap=SHOW_APV, src=lambda k: f'MOD(SMALL({K2},{k}),{M})')
    fl.add(title=dict(A='"③ 等老板审批"', B='"计划在本月及以前"', C='"打印看【待审批付款单】"'),
           head=dict(A='"单号"', B='"收款单位"', C='"付款内容"', D='"申请日期"', E='"计划付款日期"', F='"申请金额"',
                     G='"现在欠这家"', H='"欠票"', I='"申请人"', J='"提示"'),
           item=dict(common, F=lambda n, k: P('申请金额', n), G=lambda n, k: neg(P('现在欠款', n), '预付'),
                     H=lambda n, k: neg(P('欠票', n), '多收票'), I=lambda n, k: P('申请人', n),
                     J=lambda n, k: f'MID({plan_hint(n)}&{dup(n)}&{pend(n)},2,200)'),
           tot=dict(A='"合计"', F=sc('审申请')),
           trail=dict(A=f'IF({sc("审n")}=0,"（没有）","共 "&{sc("审n")}&" 单")',
                      C=f'IF({sc("审n")}>{SHOW_APV},"只列了前 {SHOW_APV} 单（合计是全部的）","")'),
           n_all=sc('审n'), cap=SHOW_APV, src=lambda k: f'MOD(SMALL({K3},{k}),{M})')
    fl.add(title=dict(A='"④ 暂缓的"', B='"全部"', C='"老板说先放一放的"'),
           head=dict(A='"单号"', B='"收款单位"', C='"付款内容"', D='"申请日期"', E='"计划付款日期"', F='"申请金额"',
                     I='"申请人"', J='"审批意见"'),
           item=dict(common, F=lambda n, k: P('申请金额', n), I=lambda n, k: P('申请人', n), J=lambda n, k: P('审批意见', n)),
           tot=dict(A='"合计"', F=sc('缓申请')),
           trail=dict(A=f'IF({sc("缓n")}=0,"（没有）","共 "&{sc("缓n")}&" 单")',
                      C=f'IF({sc("缓n")}>{SHOW_HOLD},"只列了前 {SHOW_HOLD} 单（合计是全部的）","")'),
           n_all=sc('缓n'), cap=SHOW_HOLD, src=lambda k: f'MOD(SMALL({K4},{k}),{M})', spacer=0)
    fl.build(fmts=dict(D=DATE, E=DATE, F=MONEY, G=MONEY, H=MONEY),
             aligns=dict(A=AL, B=AL, C=AL, D=AC, E=AC, F=AR, G=AR, H=AR, I=AC, J=AC), wrap='CJ',
             height=26)                        # 固定两行高：换了月份 Excel 不会自己调行高
    sc.set('尾', fl.end)
    sc.set('末行', f'{FP_R0}+{sc("尾")}-1')
    T, S0, r0 = fl.t(), fl.s(), FP_R0
    red = dict(font=Font(bold=True, color='FFC00000'), border=CF_BD)
    fl.cf(LAST, [
        ('J', 'J', lambda c: f'AND({T}=5,OR({S0}=1,{S0}=2,{S0}=3),ISNUMBER(FIND("过了",$J{r0})))', red),
        ('J', 'J', lambda c: f'AND({T}=5,$J{r0}="已付清")', dict(font=Font(color='FF548235'), border=CF_BD)),
        ('H', 'H', lambda c: f'AND({T}=5,{S0}=1,N($H{r0})>0.005)', dict(font=Font(bold=True), border=CF_BD)),
        ('G', 'H', lambda c: f'AND({T}=5,{S0}=3,ISTEXT({c}{r0}))', dict(font=Font(bold=True, color='FF2F75B5'), border=CF_BD)),
        ('C', 'C', lambda c: f'AND({T}=6,LEFT($C{r0},1)="⚠")', dict(font=Font(bold=True, color='FFC00000'))),
    ])
    hide(ws, *[CL(i) for i in range(CI('L'), CI(FIX_CUT['后月']) + 1)])
    finish(ws, LAST, sc('末行'), fl.r1, 'A4')


# ═══════════════════════════════ 待审批付款单 ═══════════════════════════════
TIP_APR = ('💡 给老板打印签字用（A4 横向一页宽）：【付款审批】里全部还没审批的单，按计划付款日期排（没填计划日期的按申请日期），'
           '「批示」「批准金额」两栏空着给老板手写；下面另列暂缓的单。老板批完，出纳把结果填回【付款审批】紫色几列（审批结果、批准金额、审批意见、审批日期）。'
           '现在欠这家（我们还欠收款单位的货款）、欠票、已批未付都是截止日的数；「预付」＝已经先付了钱，「多收票」＝对方发票开多了，'
           '「＋N笔没定价」＝还有到货没定价、没算进欠款。都批了还剩＝可用资金－本月固定支出还没付－已批未付－待审批（票据不算）。')
AP_R0 = 8


def build_aprint(ws):
    LAST = 'M'
    widths(ws, {'A': 5, 'B': 12, 'C': 12.5, 'D': 8, 'E': 14, 'F': 20, 'G': 13.5, 'H': 12.5, 'I': 13.5, 'J': 13.5, 'K': 13.5,
                'L': 16, 'M': 13.5, 'N': 10})
    title(ws, '=IF(P_公司名称="","",P_公司名称&" ")&"待审批付款单"', LAST, C_PLAN, TIP_APR)
    home_link(ws, 'N3')
    keys = ['可用', '票据', '截止月', '固没付', '已批未付', '审n', '审申请', '缓n', '缓申请', '都批了还剩', '尾', '末行']
    sc = Sc(ws, keys)
    avail, bills = acct_table(ws)
    sc.set('可用', avail)
    sc.set('票据', bills)
    sc.set('截止月', 'P_截止年月')
    fx = fix_table(ws, sc('截止月'))
    sc.set('固没付', fx['没付'])
    sc.set('已批未付', R2('SUMIFS(批_未付,批_有效,1)'))
    pi = lambda k, n: f'INDEX(批_{k},{n})'
    K1 = key_col(ws, 'BH', N_APV, lambda n: f'AND({pi("有效", n)}=1,{pi("状态码", n)}=0)', lambda n: pi('排序键', n), '待审批键')
    K2 = key_col(ws, 'BI', N_APV, lambda n: f'AND({pi("有效", n)}=1,{pi("状态码", n)}=5)', lambda n: pi('排序键', n), '暂缓键')
    sc.set('审n', f'COUNT({K1})')
    sc.set('审申请', R2('SUMIFS(批_申请金额,批_有效,1,批_状态码,0)'))
    sc.set('缓n', f'COUNT({K2})')
    sc.set('缓申请', R2('SUMIFS(批_申请金额,批_有效,1,批_状态码,5)'))
    sc.set('都批了还剩', R2(f'{sc("可用")}-{sc("固没付")}-{sc("已批未付")}-{sc("审申请")}'))

    # 第 3 行：截至日期（＝截止日期）
    ws.merge_cells('A3:B3')
    put(ws, 'A3', '截至日期', F_KPI_L, LBL, align=AC)
    put(ws, 'B3', None, F_KPI_L, LBL)
    ws.merge_cells('C3:D3')
    put(ws, 'C3', '=TEXT(P_截止,"yyyy-mm-dd")', F_SEL, FILL_AUTO, align=AC)
    put(ws, 'D3', None, F_SEL, FILL_AUTO)
    ws.merge_cells('E3:M3')
    put(ws, 'E3', (f'="（＝截止日期，首页改）　等老板审批的 "&{sc("审n")}&" 单，申请金额合计 "&TEXT({sc("审申请")},"#,##0.00")'
                   f'&IF({sc("缓n")}>0,"；另有暂缓的 "&{sc("缓n")}&" 单","")'), F_NOTE, align=AL, border=False)
    ws.row_dimensions[3].height = 24
    # 第 4、5 行：钱
    kp = [('A', 'C', '现在可用资金', sc('可用')), ('D', 'E', '手上票据（承兑）', sc('票据')),
          ('F', 'G', f'="本月（"&MOD(P_截止年月,100)&" 月）固定支出还没付"', sc('固没付')),
          ('H', 'I', '已批未付（全部）', sc('已批未付')),
          ('J', 'K', f'=IF({sc("审n")}>{SHOW_APV},"待审批合计（全部 "&{sc("审n")}&" 单）","本单待审批合计")', sc('审申请')),
          ('L', 'M', '都批了还剩\n（可用－固定没付－已批未付－待审批）', sc('都批了还剩'))]
    for c1, c2, lab, f in kp:
        ws.merge_cells(f'{c1}4:{c2}4')
        ws.merge_cells(f'{c1}5:{c2}5')
        put(ws, f'{c1}4', lab, F_KPI_L, LBL, align=ACW)
        put(ws, f'{c1}5', f'=TEXT({f},"#,##0.00;-#,##0.00;0.00")', F_KVB, FILL_TOT, align=AC)   # 文字：合并格不撑宽第一列
        for i in range(CI(c1) + 1, CI(c2) + 1):
            put(ws, f'{CL(i)}4', None, F_KPI_L, LBL)
            put(ws, f'{CL(i)}5', None, F_KVB, FILL_TOT)
    ws.row_dimensions[4].height = 44
    ws.row_dimensions[5].height = 24
    ws.conditional_formatting.add('A5:M5', FormulaRule(formula=['LEFT(A5,1)="-"'], font=Font(bold=True, color='FFC00000')))
    section(ws, 6, 'A', LAST, '="等老板审批的付款（按计划付款日期排；没填计划日期的按申请日期）　截至 "&TEXT(P_截止,"yyyy-mm-dd")', C_PLAN)
    heads = [('A', '序号'), ('B', '单号'), ('C', '申请日期'), ('D', '申请人'), ('E', '收款单位'), ('F', '付款内容'), ('G', '申请金额'),
             ('H', '计划付款日期'), ('I', '现在欠这家'), ('J', '欠票'), ('K', '这家已批未付'), ('L', '批示\n（同意/不同意/暂缓）'),
             ('M', '批准金额')]
    header(ws, 7, heads[:11], PURPLE_H, height=36)
    header(ws, 7, heads[11:], C_PLAN, height=36)

    P = lambda k, n: f'INDEX(批_{k},{n})'
    u = lambda n: P('单位号', n)
    npp = lambda n: f'INDEX(位_待定价到货,MAX(1,{u(n)}))'
    ow = lambda n: P('现在欠款', n)

    def owe(n):
        """现在欠这家：负数写「预付」；这家还有到货没定价（没算进欠款）→ 后面加「＋N笔没定价」"""
        return (f'IF(AND({ow(n)}<>"",{u(n)}>0),IF({npp(n)}>0,IF({ow(n)}<-0.005,"预付 "&TEXT(-{ow(n)},"#,##0.00"),'
                f'TEXT({ow(n)},"#,##0.00"))&"＋"&{npp(n)}&"笔没定价",{neg(ow(n), "预付")}),{neg(ow(n), "预付")})')
    # 只列前几条时，合计只算列出来的（老板签字的单子上看得到的才合计）
    shown = lambda K, cap, n_all, all_: (f'IF({n_all}>{cap},ROUND(SUMIFS(批_申请金额,批_有效,1,{K},"<="&SMALL({K},{cap})),2),'
                                         f'{all_})')
    item = dict(A=lambda n, k: k, B=lambda n, k: P('单号', n), C=lambda n, k: dt0(P('申请日期', n)), D=lambda n, k: P('申请人', n),
                E=lambda n, k: P('单位', n), F=lambda n, k: P('付款内容', n), G=lambda n, k: P('申请金额', n),
                H=lambda n, k: dt0(P('计划付款日期', n)), I=lambda n, k: owe(n),
                J=lambda n, k: neg(P('欠票', n), '多收票'), K=lambda n, k: f'IF({u(n)}=0,"",INDEX(位_已批未付,MAX(1,{u(n)})))')
    fl = Flow(ws, sc, AP_R0, 'ABCDEFGHIJKLM')
    fl.add(item=item, tot=dict(B='"合计"', G=shown(K1, SHOW_APV, sc('审n'), sc('审申请'))),
           trail=dict(B=f'IF({sc("审n")}=0,"（没有）","共 "&{sc("审n")}&" 单")',
                      F=(f'IF({sc("审n")}=0,"现在没有等审批的",IF({sc("审n")}>{SHOW_APV},'
                         f'"只列了前 {SHOW_APV} 单（合计只算这些；全部见顶上）",""))')),
           n_all=sc('审n'), cap=SHOW_APV, src=lambda k: f'MOD(SMALL({K1},{k}),{M})')
    fl.add(title=dict(B='"暂缓的"', C='"（全部）"', F='"老板说先放一放的，也可以再批"'),
           head=dict({c: f'"{t}"' for c, t in heads if '\n' not in t}, L='"批示"'),
           item=item, tot=dict(B='"合计"', G=shown(K2, SHOW_HOLD, sc('缓n'), sc('缓申请'))),
           trail=dict(B=f'IF({sc("缓n")}=0,"（没有）","共 "&{sc("缓n")}&" 单")',
                      F=f'IF({sc("缓n")}>{SHOW_HOLD},"只列了前 {SHOW_HOLD} 单（合计只算这些）","")'),
           extra=[{}, dict(I='"老板签字："', L='"日期："')],
           n_all=sc('缓n'), cap=SHOW_HOLD, src=lambda k: f'MOD(SMALL({K2},{k}),{M})', spacer=0)
    fl.build(fmts=dict(C=DATE, G=MONEY, H=DATE, I=MONEY, J=MONEY, K=MONEY),
             aligns=dict(A=AC, B=AL, C=AC, D=AC, E=AL, F=AL, G=AR, H=AC, I=AR, J=AR, K=AR, L=AC, M=AR), height=24,
             wrap='EFI', shrink='KL')                  # 收款单位、付款内容、现在欠这家（文字）换行；金额列还是缩小
    sc.set('尾', fl.end)
    sc.set('末行', f'{AP_R0}+{sc("尾")}-1')
    T, r0 = fl.t(), AP_R0
    fl.cf(LAST, [
        ('I', 'J', lambda c: f'AND({T}=5,ISTEXT({c}{r0}))', dict(font=Font(bold=True, color='FF2F75B5'), border=CF_BD)),
        ('J', 'K', lambda c: f'{T}=71', dict(border=SIGN_BD)),
        ('M', 'M', lambda c: f'{T}=71', dict(border=SIGN_BD)),
        ('I', 'I', lambda c: f'{T}=71', dict(font=Font(bold=True))),
        ('L', 'L', lambda c: f'{T}=71', dict(font=Font(bold=True))),
    ])
    hide(ws, *[CL(i) for i in range(CI('O'), CI('BI') + 1)])
    finish(ws, LAST, sc('末行'), fl.r1, f'A{AP_R0}', hdr_rows='7:7')


# ═══════════════════════════════ 采购计划汇总 ═══════════════════════════════
TIP_PP = ('💡 选日期（空着＝截止日）：① 这天提交的采购计划，打印就是「今日采购计划」。'
          '选月份（2026-10、202610 或只填 10；空着＝截止日那个月）：② 这个月按状态、按提交人汇总。'
          '③ 等老板审批的、④ 同意了还没买的（全部，按需用日期排，到了日子标红）。打印选「当天的单」只印 ①，选「全部」印整张。')
PP_R0 = 7
PER_C = dict(姓名='AR', 第一次='AS', 条数='AT', 金额='AU', 计数='AV')


def build_ppsum(ws):
    LAST = 'J'
    widths(ws, {'A': 6, 'B': 10, 'C': 24, 'D': 9, 'E': 7, 'F': 13.5, 'G': 14, 'H': 12.5, 'I': 12.5, 'J': 12, 'K': 10})
    title(ws, '=IF(P_公司名称="","",P_公司名称&"　")&"采购计划汇总"', LAST, C_PLAN, TIP_PP)
    selector(ws, 'A3', '选日期', 'C3', None, fmt=DATE)
    ws.merge_cells('A3:B3')
    dv_date(ws, 'C3')
    selector(ws, 'D3', '选月份', 'F3', None, fmt=YM_FMT)
    ws.merge_cells('D3:E3')
    selector(ws, 'G3', '打印', 'H3', None, dv_formula='"当天的单,全部"', prompt='当天的单＝只印 ①（今日采购计划）；全部＝整张；空着＝当天的单')
    home_link(ws, 'K3')                       # 放在打印范围（A～J）外面
    ws.row_dimensions[3].height = 30
    keys = ['d', '认出d', '输入', '认出前', 'm', '认出', 'Y', 'mo', '全部印', '①前', '①n', '①金额', '月n', '月金额', '人数',
            '各人n', '各人金额', '其他n', '其他金额', '③n', '③金额', '④n', '④金额', '尾', '①末行', '全末行', '末行', '最晚']
    sc = Sc(ws, keys)
    d = sc('d')
    sc.set('d', 'IF(AND(ISNUMBER(C3),C3>=36526,C3<73051),INT(C3),P_截止)')
    sc.set('认出d', 'IF(TRIM(C3&"")="",1,IF(AND(ISNUMBER(C3),C3>=36526,C3<73051),1,0))')
    # 采购计划里最晚的提交日期（不超过日期上限）：比截止日晚（资金台帐还没记到那天）时提醒去选
    sc.set('最晚', f'IFERROR(INT(LARGE(购_排序键,COUNTIF(购_排序键,">="&(P_日期上限+1)*{M})+1)/{M}),0)')
    month_selector(ws, sc, 'F3')
    m, Y, mo = sc('m'), sc('Y'), sc('mo')
    sc.set('全部印', 'IF(TRIM(H3&"")="全部",1,0)')
    sc.set('①前', f'COUNTIF(购_排序键,"<"&{d}*{M})')
    sc.set('①n', f'COUNTIFS(购_排序键,">="&{d}*{M},购_排序键,"<"&({d}+1)*{M})')
    sc.set('①金额', R2(f'SUMIFS(购_预计金额,购_有效,1,购_日期,{d})'))
    inm = f'购_有效,1,购_年月,{m}'
    sc.set('月n', f'COUNTIFS({inm})')
    sc.set('月金额', R2(f'SUMIFS(购_预计金额,{inm})'))
    # 人员逐个（人员_姓名 顺序，同名只算第一个）
    C = PER_C
    for k, c in C.items():
        ws[f'{c}2'] = k
        ws[f'{c}2'].font = F_HELP
    for i in range(N_PER):
        n, r = i + 1, 3 + i
        nm, fs = f'{C["姓名"]}{r}', f'{C["第一次"]}{r}'
        hset(ws, nm, f'INDEX(人员_姓名,{n})&""')
        hset(ws, fs, f'IF({nm}="",0,IF(IFERROR(MATCH({esc(nm)},人员_姓名,0),0)={n},1,0))')
        hset(ws, f'{C["条数"]}{r}', f'IF({fs}<>1,0,COUNTIFS({inm},购_提交人,{esc(nm)}))')
        hset(ws, f'{C["金额"]}{r}', f'IF({fs}<>1,0,ROUND(SUMIFS(购_预计金额,{inm},购_提交人,{esc(nm)}),2))')
    counter(ws, C['计数'], 3, N_PER, lambda i: f'${C["第一次"]}${3 + i}=1')
    pr = lambda k: f'${C[k]}$3:${C[k]}${N_PER + 2}'
    sc.set('人数', cnt(C['计数'], 3, N_PER))
    sc.set('各人n', f'SUM({pr("条数")})')
    sc.set('各人金额', R2(f'SUM({pr("金额")})'))
    sc.set('其他n', f'{sc("月n")}-{sc("各人n")}')
    sc.set('其他金额', R2(f'{sc("月金额")}-{sc("各人金额")}'))
    # ③④ 排序辅助：需用日期（没填＝99999）×SORT_M＋n
    g = lambda k, n: f'INDEX(购_{k},{n})'
    need_key = lambda n: f'IF({g("需用日期", n)}>0,{g("需用日期", n)},99999)*{M}+{n}'
    K3 = key_col(ws, 'BH', N_PP, lambda n: f'AND({g("有效", n)}=1,{g("状态码", n)}=0)', need_key, '③键')
    K4 = key_col(ws, 'BI', N_PP, lambda n: f'AND({g("有效", n)}=1,{g("状态码", n)}=1)', need_key, '④键')
    sc.set('③n', f'COUNT({K3})')
    sc.set('③金额', R2('SUMIFS(购_预计金额,购_有效,1,购_状态码,0)'))
    sc.set('④n', f'COUNT({K4})')
    sc.set('④金额', R2('SUMIFS(购_预计金额,购_有效,1,购_状态码,1)'))

    ws.merge_cells('A4:J4')
    put(ws, 'A4', (f'="实际用的：日期 "&TEXT({d},"yyyy-mm-dd")&IF(TRIM(C3&"")="","（空着＝截止日）","")'
                   f'&IF({sc("认出d")}=0,"（⚠ 没认出这个日期，先按截止日）","")'
                   f'&IF(AND(TRIM(C3&"")="",{sc("最晚")}>{d}),"（⚠ 采购计划最晚有 "&TEXT({sc("最晚")},"yyyy-mm-dd")'
                   f'&" 提交的，比截止日晚：要印那天的单，黄格填那天）","")'
                   f'&"；月份 "&{Y}&" 年 "&{mo}&" 月"&IF(TRIM(F3&"")="","（空着＝截止日那个月）","")'
                   f'&IF({sc("认出")}=0,"（⚠ 没认出这个月份，先按截止日那个月）","")'
                   f'&"；打印："&IF({sc("全部印")}=1,"整张","只印 ①（当天的单）")'), F_NOTE, align=ALW, border=False)
    ws.row_dimensions[4].height = 28
    ws.conditional_formatting.add('A4', FormulaRule(formula=['ISNUMBER(FIND("⚠",$A$4))'],
                                                    font=Font(bold=True, color='FFC00000')))
    section(ws, 5, 'A', LAST, f'="① "&TEXT({d},"yyyy-mm-dd")&" 提交的采购计划（今日采购计划）"', C_PLAN)
    header(ws, 6, [('A', '序号'), ('B', '提交人'), ('C', '物料及规格'), ('D', '数量'), ('E', '单位'), ('F', '预计金额'),
                   ('G', '建议供应商'), ('H', '需用日期'), ('I', '审批'), ('J', '状态')], PURPLE_H, height=30)

    fl = Flow(ws, sc, PP_R0, 'ABCDEFGHIJ')
    base = dict(A=lambda n, k: k, B=lambda n, k: g('提交人', n), C=lambda n, k: g('物料', n),
                D=lambda n, k: f'IF({g("数量", n)}=0,"",{g("数量", n)})', E=lambda n, k: g('计量单位', n),
                F=lambda n, k: g('预计金额', n), G=lambda n, k: g('建议供应商', n), H=lambda n, k: dt0(g('需用日期', n)))
    fl.add(item=dict(base, I=lambda n, k: g('审批', n), J=lambda n, k: g('状态', n)),
           tot=dict(B='"合计"', F=sc('①金额')),
           trail=dict(B=f'IF({sc("①n")}=0,"（没有）","共 "&{sc("①n")}&" 条")',
                      C=f'IF({sc("①n")}=0,"这天没有提交采购计划",IF({sc("①n")}>{SHOW_PP},"只列了前 {SHOW_PP} 条（合计是全部的）",""))'),
           n_all=sc('①n'), cap=SHOW_PP, src=lambda k: f'MOD(SMALL(购_排序键,{sc("①前")}+{k}),{M})')
    st_name = lambda n: 'CHOOSE(' + n + ',' + ','.join(f'"{x}"' for x in ST_PP) + ')'
    st_note = lambda n: 'CHOOSE(' + n + ',' + ','.join(f'"{x}"' for x in ST_PP_NOTE) + ')'
    fl.add(title=dict(A='"②"', B='"本月汇总"', C=f'{Y}&" 年 "&{mo}&" 月提交的 · 按状态"'),
           head=dict(B='"状态"', C='"说明"', D='"条数"', F='"预计金额"'),
           item=dict(B=lambda n, k: st_name(n), C=lambda n, k: st_note(n),
                     D=lambda n, k: f'COUNTIFS({inm},购_状态码,{n}-1)',
                     F=lambda n, k: R2(f'SUMIFS(购_预计金额,{inm},购_状态码,{n}-1)')),
           tot=dict(B='"合计"', D=sc('月n'), F=sc('月金额')),
           n_all='7', cap=7, src=lambda k: k)
    other = 999
    PT = lambda k, n: f'INDEX({pr(k)},{n})'
    fl.add(title=dict(B='"按提交人"', C=f'{Y}&" 年 "&{mo}&" 月提交的"'),
           head=dict(B='"提交人"', D='"条数"', F='"预计金额"'),
           item=dict(B=lambda n, k: f'IF({n}={other},"其他",{PT("姓名", n)})',
                     C=lambda n, k: f'IF({n}={other},"不在【基础资料】人员里的、没填的","")',
                     D=lambda n, k: f'IF({n}={other},{sc("其他n")},{PT("条数", n)})',
                     F=lambda n, k: f'IF({n}={other},{sc("其他金额")},{PT("金额", n)})'),
           tot=dict(B='"合计"', D=sc('月n'), F=sc('月金额')),
           n_all=f'({sc("人数")}+1)', cap=N_PER + 1,
           src=lambda k: f'IF({k}>{sc("人数")},{other},{kth(k, C["计数"], 3, N_PER)})')
    late = lambda n: f'IF(AND({g("需用日期", n)}>0,{g("需用日期", n)}<=P_截止),"已到日子","")'
    for tag, nm, sub, key, nn, amt, sp in (('③', '③ 待审批', '老板还没批的', K3, '③n', '③金额', 1),
                                           ('④', '④ 同意没买', '老板同意了、还没下单的', K4, '④n', '④金额', 0)):
        fl.add(title=dict(A=f'"{tag}"', B=f'"{nm[2:]}"', C=f'"{sub}（全部，按需用日期排）"'),
               head=dict(A='"序号"', B='"提交人"', C='"物料"', D='"数量"', E='"单位"', F='"预计金额"', G='"建议供应商"',
                         H='"需用日期"', I='"提交日期"', J='"提示"'),
               item=dict(base, I=lambda n, k: g('日期', n), J=lambda n, k: late(n)),
               tot=dict(B='"合计"', F=sc(amt)),
               trail=dict(B=f'IF({sc(nn)}=0,"（没有）","共 "&{sc(nn)}&" 条")',
                          C=f'IF({sc(nn)}>{SHOW_PP},"只列了前 {SHOW_PP} 条（合计是全部的）","")'),
               n_all=sc(nn), cap=SHOW_PP, src=lambda k, key=key: f'MOD(SMALL({key},{k}),{M})', spacer=sp)
    fl.build(fmts=dict(A=NUM, D=NUM, F=MONEY, H=DATE, I=DATE),
             aligns=dict(A=AC, B=AL, C=AL, D=AR, E=AC, F=AR, G=AL, H=AC, I=AC, J=AC), shrink='CG')
    sc.set('尾', fl.end)
    sc.set('①末行', f'{fl.row_of(0, 5)}')
    sc.set('全末行', f'{PP_R0}+{sc("尾")}-1')
    sc.set('末行', f'IF({sc("全部印")}=1,{sc("全末行")},{sc("①末行")})')
    T, S0, r0 = fl.t(), fl.s(), PP_R0
    red = dict(font=Font(bold=True, color='FFC00000'), border=CF_BD)
    fl.cf(LAST, [
        ('H', 'H', lambda c: f'AND({T}=5,$J{r0}="已到日子")', red),
        ('J', 'J', lambda c: f'AND({T}=5,$J{r0}="已到日子")', red),
        ('B', 'J', lambda c: f'AND({T}=5,OR({S0}=2,{S0}=3),$B{r0}="其他")', dict(font=Font(color='FF808080'), border=CF_BD)),
    ])
    hide(ws, *[CL(i) for i in range(CI('L'), CI('BI') + 1)])
    finish(ws, LAST, sc('末行'), fl.r1, 'A5')          # 只冻标题、选择格、实际用的（① 的表头跟下面几段对不上）


def build(wb, ctx=None):
    build_fplan(wb[SH_FPLAN])
    build_aprint(wb[SH_APRINT])
    build_ppsum(wb[SH_PPSUM])
